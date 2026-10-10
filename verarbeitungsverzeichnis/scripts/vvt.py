#!/usr/bin/env python3
"""Lokales VVT mit nachvollziehbaren Änderungen; keine Rechtsentscheidung oder Synchronisation."""
from __future__ import annotations
import argparse
import copy
from contextlib import contextmanager
from datetime import date, datetime, timezone
import hashlib
import html
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
import uuid
import xml.etree.ElementTree as ET
import zipfile

SCHEMA = 1
MAX_BYTES = 25 * 1024 * 1024
MAX_ACTIVITIES = 5000  # technische Schutzgrenze, keine rechtliche Größenfreistellung
TEXT_FIELDS = ("title", "owner", "purpose", "joint_controller", "controller_contacts",
               "processing_categories", "data_subjects", "data_categories", "recipients",
               "transfers", "retention", "toms", "legal_basis", "systems", "processor_contracts",
               "sources", "next_review")
TRIGGERS = ("high_risk", "art35_3", "positive_list")
SCREEN_DEFAULTS = {**{k: "offen" for k in TRIGGERS}, "severity": "offen", "likelihood": "offen",
                   "priority": "offen", "criteria": "", "rationale": "", "sources": ""}
REVIEW_DEFAULTS = {"decision": "offen", "reviewer": "", "reviewed_at": "", "rationale": "",
                   "sources": "", "basis_sha256": ""}
LABELS = {"id": "VT-ID", "activity_revision": "Tätigkeitsrevision", "title": "Bezeichnung",
 "owner": "Verantwortliche Person", "role": "Rolle", "status": "Bearbeitungsstatus",
 "purpose": "Zwecke", "joint_controller": "Gemeinsam Verantwortliche",
 "controller_contacts": "Verantwortliche Auftraggeber und Kontaktdaten",
 "processing_categories": "Verarbeitungskategorien je Auftraggeber",
 "data_subjects": "Kategorien betroffener Personen", "data_categories": "Datenkategorien",
 "recipients": "Empfängerkategorien", "transfers": "Drittlandtransfers und Nachweise",
 "retention": "Löschung: Fristen, Auslöser und Begründung", "toms": "Technische und organisatorische Maßnahmen",
 "legal_basis": "Rechtsgrundlagen und Prüfung", "systems": "Systeme und Dienstleister",
 "processor_contracts": "Auftragsverarbeitungsverträge", "sources": "Belegverweise und Stand",
 "next_review": "Nächste Überprüfung", "high_risk": "Voraussichtlich hohes Risiko",
 "art35_3": "Tatbestand Artikel 35 Absatz 3", "positive_list": "Einschlägige DSFA-Pflichtliste",
 "severity": "Schwere (interne Einordnung)", "likelihood": "Wahrscheinlichkeit (interne Einordnung)",
 "priority": "Interne Bearbeitungspriorität", "criteria": "Prüfkriterien und Fallbezug",
 "rationale": "Begründung", "decision": "DSFA-Entscheidung", "reviewer": "Prüfende Person",
 "reviewed_at": "Prüfdatum", "basis_sha256": "Geprüfter Sachstand (SHA-256)"}
DECISION_LABELS = {"offen": "Noch nicht entschieden", "erforderlich": "DSFA erforderlich", "begruendet_nicht_erforderlich": "DSFA nach begründeter Prüfung nicht erforderlich"}
SHEETS = {
 "Taetigkeiten": ("id", "activity_revision", "title", "role", "owner", "status", "next_review", "purpose", "systems", "legal_basis", "sources"),
 "Art30": ("id", "joint_controller", "controller_contacts", "processing_categories", "data_subjects", "data_categories", "recipients", "transfers", "retention", "toms", "processor_contracts"),
 "Screening": ("id", *SCREEN_DEFAULTS),
}

class VVTError(ValueError):
    pass

def utcnow():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()

def canonical(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n").encode("utf-8")

def digest(value):
    return hashlib.sha256(canonical(value)).hexdigest()

def fact_hash(activity):
    return digest({k: v for k, v in activity.items() if k not in ("review", "activity_revision")})

def review_state(activity):
    r = activity["review"]
    if r["decision"] == "offen":
        return "offen"
    return "aktuell" if r["basis_sha256"] == fact_hash(activity) else "veraltet"

def reject_duplicate_keys(pairs):
    out = {}
    for key, value in pairs:
        if key in out:
            raise VVTError(f"Doppelter JSON-Schlüssel: {key}")
        out[key] = value
    return out

def read_json(path):
    raw = Path(path).read_bytes()
    if len(raw) > MAX_BYTES:
        raise VVTError("Datei überschreitet die technische Grenze von 25 MiB.")
    return json.loads(raw.decode("utf-8-sig"), object_pairs_hook=reject_duplicate_keys,
                      parse_constant=lambda s: (_ for _ in ()).throw(VVTError("Ungültige JSON-Zahl.")))

def human(name):
    if not isinstance(name, str) or len(name.strip().split()) < 2:
        raise VVTError("Bitte Vor- und Nachnamen der verantwortlichen Person angeben.")
    if re.search(r"\b(bot|chatgpt|claude|codex|ki|agent|system|automatik)\b", name, re.I):
        raise VVTError("Eine Maschinenbezeichnung ersetzt keine namentlich verantwortliche Person.")

def nonempty(value, label):
    if not isinstance(value, str) or not value.strip():
        raise VVTError(f"{label} fehlt.")

def empty_activity(identifier):
    return {"id": identifier, "activity_revision": 0, "role": "controller", "status": "geplant",
            **{k: "" for k in TEXT_FIELDS}, "screening": copy.deepcopy(SCREEN_DEFAULTS),
            "review": copy.deepcopy(REVIEW_DEFAULTS)}

def validate(doc):
    if not isinstance(doc, dict) or set(doc) != {"schema_version", "register_id", "organization", "revision", "updated_at", "activities", "history"}:
        raise VVTError("Unbekanntes oder unvollständiges Registerschema.")
    if type(doc["schema_version"]) is not int or doc["schema_version"] != SCHEMA or type(doc["revision"]) is not int or doc["revision"] < 0:
        raise VVTError("Ungültige Schema- oder Registerrevision.")
    try:
        uuid.UUID(doc["register_id"])
    except (ValueError, TypeError, AttributeError) as exc:
        raise VVTError("Ungültige Register-ID.") from exc
    org = doc["organization"]
    if not isinstance(org, dict) or set(org) != {"name", "contact", "representative", "dpo"} or any(not isinstance(v, str) for v in org.values()):
        raise VVTError("Organisation muss name, contact, representative und dpo als Text enthalten.")
    nonempty(org["name"], "Organisation")
    if not isinstance(doc["activities"], list) or len(doc["activities"]) > MAX_ACTIVITIES:
        raise VVTError("Tätigkeitsliste ungültig oder technische Grenze überschritten.")
    if not isinstance(doc["history"], list) or not isinstance(doc["updated_at"], str):
        raise VVTError("Änderungsprotokoll oder Datum ungültig.")
    def check_texts(value):
        if isinstance(value, dict):
            for k, v in value.items():
                check_texts(k)
                check_texts(v)
        elif isinstance(value, list):
            for item in value:
                check_texts(item)
        elif isinstance(value, str):
            if len(value) > 30000 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", value):
                raise VVTError("Text enthält unzulässige Steuerzeichen oder überschreitet 30000 Zeichen.")
    check_texts(doc)
    seen = set()
    for a in doc["activities"]:
        if not isinstance(a, dict) or set(a) != set(empty_activity("VT-0001")):
            raise VVTError("Unbekannte oder fehlende Tätigkeitsfelder.")
        if not isinstance(a["id"], str) or not re.fullmatch(r"VT-[0-9]{4,8}", a["id"]) or a["id"] in seen:
            raise VVTError("Ungültige oder doppelte VT-ID.")
        seen.add(a["id"])
        if type(a["activity_revision"]) is not int or a["activity_revision"] < 1:
            raise VVTError(f"{a['id']}: ungültige Tätigkeitsrevision.")
        if a["role"] not in ("controller", "processor") or a["status"] not in ("geplant", "aktiv", "pausiert", "beendet"):
            raise VVTError(f"{a['id']}: ungültige Rolle oder Status.")
        for field in TEXT_FIELDS:
            if not isinstance(a[field], str) or len(a[field]) > 30000:
                raise VVTError(f"{a['id']}: {field} muss Text mit höchstens 30000 Zeichen sein.")
        nonempty(a["title"], f"{a['id']}: Bezeichnung")
        if a["next_review"]:
            try:
                date.fromisoformat(a["next_review"])
            except ValueError as exc:
                raise VVTError("Prüfdatum muss YYYY-MM-DD oder leer sein.") from exc
        s, r = a["screening"], a["review"]
        if not isinstance(s, dict) or set(s) != set(SCREEN_DEFAULTS) or any(not isinstance(v, str) for v in s.values()):
            raise VVTError("Screeningfelder fehlen oder sind unbekannt.")
        if any(s[k] not in ("offen", "ja", "nein") for k in TRIGGERS):
            raise VVTError("Screeningtrigger müssen offen, ja oder nein sein.")
        if any(s[k] not in ("offen", "gering", "mittel", "hoch") for k in ("severity", "likelihood")) or s["priority"] not in ("offen", "1", "2", "3"):
            raise VVTError("Ungültige interne Risikoeinordnung.")
        if not isinstance(r, dict) or set(r) != set(REVIEW_DEFAULTS) or any(not isinstance(v, str) for v in r.values()):
            raise VVTError("Prüfungsfelder fehlen oder sind unbekannt.")
        if r["decision"] not in ("offen", "erforderlich", "begruendet_nicht_erforderlich"):
            raise VVTError("Unbekannte DSFA-Entscheidung.")
        if r["decision"] != "offen":
            human(r["reviewer"])
            for field in ("reviewed_at", "rationale", "sources", "basis_sha256"):
                nonempty(r[field], "Prüfung " + field)
            try:
                stamp = datetime.fromisoformat(r["reviewed_at"])
                if stamp.tzinfo is None:
                    raise ValueError()
            except ValueError as exc:
                raise VVTError("Prüfdatum muss ein ISO-Zeitstempel mit Zeitzone sein.") from exc
            if not re.fullmatch(r"[0-9a-f]{64}", r["basis_sha256"]):
                raise VVTError("Ungültiger Prüfhash.")
            if review_state(a) == "aktuell" and r["decision"] == "begruendet_nicht_erforderlich" and any(s[k] != "nein" for k in TRIGGERS):
                raise VVTError("Keine negative DSFA-Entscheidung bei offenem oder positivem Pflichttrigger.")
    return doc

def read_register(path):
    return validate(read_json(path))

def findings(doc):
    result = []
    for a in doc["activities"]:
        # Diese Lückenliste unterscheidet gesetzliche Rollenfelder und interne Ergänzungen.
        required = ("owner", "transfers", "toms") + (("purpose", "data_subjects", "data_categories", "recipients", "retention") if a["role"] == "controller" else ("controller_contacts", "processing_categories"))
        missing = [LABELS[k] for k in required if not a[k].strip() or a[k].strip().lower() == "unbekannt"]
        result.append({"id": a["id"], "title": a["title"], "role": a["role"], "gaps": missing,
                       "screening_open": [k for k in TRIGGERS if a["screening"][k] == "offen"],
                       "mandatory_trigger": any(a["screening"][k] == "ja" for k in TRIGGERS),
                       "review_state": review_state(a), "decision": a["review"]["decision"],
                       "next_review": a["next_review"],
                       "overdue": bool(a["next_review"] and a["next_review"] <= date.today().isoformat())})
    return result

def atomic_bytes(path, data, exclusive=False):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if exclusive and path.exists():
        raise VVTError(f"Datei existiert bereits: {path}")
    fd, temp = tempfile.mkstemp(prefix="." + path.name + ".", dir=path.parent)
    try:
        os.fchmod(fd, 0o600)
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        if exclusive:
            os.link(temp, path)  # atomarer Ausschluss vorhandener Ziele
        else:
            os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)

@contextmanager
def locked(path):
    lock = Path(str(path) + ".lock")
    try:
        fd = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    except FileExistsError as exc:
        raise VVTError("Register ist gesperrt. Laufenden Vorgang prüfen; verwaiste .lock-Datei nur nach manueller Prüfung entfernen.") from exc
    try:
        os.write(fd, f"PID {os.getpid()} {utcnow()}".encode())
        os.close(fd)
        yield
    finally:
        lock.unlink(missing_ok=True)

def create(path, organization, actor, reason):
    human(actor)
    nonempty(reason, "Änderungsanlass")
    now = utcnow()
    doc = {"schema_version": SCHEMA, "register_id": str(uuid.uuid4()), "organization": organization,
           "revision": 0, "updated_at": now, "activities": [], "history": [
               {"revision": 0, "at": now, "actor": actor, "reason": reason, "operation": "init", "changed_ids": []}]}
    validate(doc)
    atomic_bytes(path, canonical(doc), exclusive=True)
    return doc

def mutate(path, expected_revision, expected_sha256, actor, reason, operation, transform):
    human(actor)
    nonempty(reason, "Änderungsanlass")
    with locked(path):
        before = read_register(path)
        if before["revision"] != expected_revision or digest(before) != expected_sha256:
            raise VVTError("Versionskonflikt: Revision oder SHA-256 stimmt nicht. Neu exportieren und Änderungen abgleichen.")
        after = copy.deepcopy(before)
        changed = transform(after)
        validate(after)
        if not changed:
            return before  # unveränderter Roundtrip erzeugt keine Scheinrevision
        after["revision"] += 1
        after["updated_at"] = utcnow()
        after["history"].append({"revision": after["revision"], "at": after["updated_at"], "actor": actor,
                                 "reason": reason, "operation": operation, "changed_ids": sorted(changed),
                                 "previous_sha256": digest(before)})
        validate(after)
        archive = Path(path).parent / (Path(path).name + ".history") / f"r{before['revision']:06d}-{digest(before)}.json"
        if not archive.exists():
            atomic_bytes(archive, canonical(before), exclusive=True)
        atomic_bytes(path, canonical(after))
        return after

def merge_activities(doc, patches):
    if not isinstance(patches, list):
        raise VVTError("Tätigkeitsänderungen müssen eine Liste sein.")
    current = {a["id"]: a for a in doc["activities"]}
    seen, changed = set(), []
    for patch in patches:
        if not isinstance(patch, dict) or "id" not in patch or patch["id"] in seen:
            raise VVTError("Tätigkeitsänderung ohne eindeutige VT-ID.")
        identifier = patch["id"]
        seen.add(identifier)
        unknown = set(patch) - (set(empty_activity(identifier)) - {"review"})
        if unknown:
            raise VVTError("Nicht bearbeitbare Felder: " + ", ".join(sorted(unknown)))
        old = current.get(identifier)
        if old and patch.get("activity_revision", old["activity_revision"]) != old["activity_revision"]:
            raise VVTError(f"{identifier}: Tätigkeitsrevision widerspricht dem Register.")
        if not old and patch.get("activity_revision", 0) not in (0, None, ""):
            raise VVTError(f"{identifier}: neue Tätigkeit muss mit Revision 0 beginnen.")
        new = copy.deepcopy(old or empty_activity(identifier))
        for key, value in patch.items():
            if key == "screening":
                if not isinstance(value, dict) or set(value) - set(SCREEN_DEFAULTS):
                    raise VVTError("Unbekannte Screeningfelder.")
                new[key].update(value)
            elif key != "activity_revision":
                new[key] = value
        if not old or fact_hash(new) != fact_hash(old):
            new["activity_revision"] = (old["activity_revision"] if old else 0) + 1
            current[identifier] = new
            changed.append(identifier)
    doc["activities"] = list(current.values())
    return changed

def upsert(path, patches, revision, sha256, actor, reason):
    return mutate(path, revision, sha256, actor, reason, "upsert", lambda doc: merge_activities(doc, patches))

def review(path, identifier, decision, rationale, sources, revision, sha256, actor, reason):
    if decision not in ("erforderlich", "begruendet_nicht_erforderlich"):
        raise VVTError("Prüfentscheidung muss erforderlich oder begruendet_nicht_erforderlich sein.")
    nonempty(rationale, "Fallbezogene Prüfbegründung")
    nonempty(sources, "Prüfquellen und Belege")
    def change(doc):
        a = next((a for a in doc["activities"] if a["id"] == identifier), None)
        if a is None:
            raise VVTError("VT-ID nicht vorhanden.")
        if decision == "begruendet_nicht_erforderlich" and any(a["screening"][k] != "nein" for k in TRIGGERS):
            raise VVTError("Offene oder positive Pflichttrigger sperren eine negative DSFA-Entscheidung; niedrige Priorität ändert das nicht.")
        a["review"] = {"decision": decision, "reviewer": actor, "reviewed_at": utcnow(),
                       "rationale": rationale, "sources": sources, "basis_sha256": fact_hash(a)}
        return [identifier]
    return mutate(path, revision, sha256, actor, reason, "review", change)

def excel_text(cell, value):
    # Explizite Stringzellen: auch =HYPERLINK(...), +, -, @ bleiben bloße Inhalte.
    cell.value = str(value)
    cell.data_type = "s"
    cell.number_format = "@"

def update_organization(path, patch, revision, sha256, actor, reason):
    if not isinstance(patch, dict) or set(patch) - {"name", "contact", "representative", "dpo"}:
        raise VVTError("Unbekannte Organisationsfelder.")
    def change(doc):
        updated = {**doc["organization"], **patch}
        if updated == doc["organization"]:
            return []
        doc["organization"] = updated
        # Geänderter Verantwortlichenkontext verlangt neue Prüfung; alte Entscheidung
        # bleibt in der unveränderten Vorfassung erhalten.
        for a in doc["activities"]:
            a["review"] = copy.deepcopy(REVIEW_DEFAULTS)
            a["activity_revision"] += 1
        return [a["id"] for a in doc["activities"]] + ["organisation"]
    return mutate(path, revision, sha256, actor, reason, "update-org", change)

def export_xlsx(doc, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    overview = wb.active
    overview.title = "Uebersicht"
    instructions = [
      ("Verarbeitungsverzeichnis", doc["organization"]["name"]),
      ("Dateistand", f"Registerrevision {doc['revision']} | {doc['updated_at']}"),
      ("Bearbeitung", "Taetigkeiten, Art30 und Screening sind editierbar. VT-ID und Tätigkeitsrevision unverändert lassen."),
      ("Neue Tätigkeit", "Neue eindeutige VT-ID mit Revision 0 in Taetigkeiten anlegen; fehlende Zusatzblätter bleiben offen."),
      ("Leer und unbekannt", "Leere Zelle = noch nicht erhoben. Unbekannt ausdrücklich als unbekannt eintragen. Nein und nicht einschlägig begründen."),
      ("Sicherung", "Import benötigt denselben Registerstand. Fehlende Zeilen löschen nichts. Beendigung nur über Status beendet."),
      ("Prüfung", "Sachänderungen entwerten frühere DSFA-Prüfungen. Interne Priorität 1–3 ist keine gesetzliche Risikoklasse."),
      ("Freigabe", "DSFA-Entscheidungen ausschließlich über review mit benannter Person, Begründung und Quellen festhalten."),
      ("Organisation", "Organisationsdaten und Prüfprotokoll sind hier nur lesbar. Organisationsänderungen über update-org; Prüfungen über review."),
      ("Grenze", "Lokale Datei ohne Cloudabgleich. Dieses Blatt aktualisiert sich erst beim erneuten Export; es berechnet keine Rechtsfreigabe."),
      ("Lange Texte", "Bei sehr langen Zellen kann die Excel-Zeilenhöhe nicht den gesamten Text zeigen. Inhalt in der Bearbeitungszeile oder der DOCX-Lesefassung lesen; er wird nicht gekürzt."),
      ("Formeln", "Eingabeblätter erlauben nur Werte; Formeln, Makros und externe Verknüpfungen werden abgewiesen."),
    ]
    for row in instructions:
        overview.append(row)
    overview.append([])
    overview.append(["VT-ID", "Bezeichnung", "Lücken", "Gültigkeit der Prüfung", "Nächste Überprüfung", "DSFA-Entscheidung"])
    for item in findings(doc):
        overview.append([item["id"], item["title"], "; ".join(item["gaps"]) or "Keine im technischen Mindestcheck", item["review_state"], item["next_review"], DECISION_LABELS[item["decision"]]])
    for name, keys in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(keys)  # stabile maschinenlesbare Feldnamen
        ws.append([LABELS.get(k, k) for k in keys])
        for a in doc["activities"]:
            values = a["screening"] if name == "Screening" else a
            row = [a["id"] if k == "id" else values.get(k, "") for k in keys]
            if "next_review" in keys and a["next_review"]:
                row[keys.index("next_review")] = date.fromisoformat(a["next_review"])
            ws.append(row)
        ws.auto_filter.ref = f"A2:{get_column_letter(len(keys))}{max(ws.max_row, 3)}"
        ws.freeze_panes = "C3" if name == "Taetigkeiten" else "B3"
        for col, key in enumerate(keys, 1):
            choices = {"role": "controller,processor", "status": "geplant,aktiv,pausiert,beendet", **{k:"offen,ja,nein" for k in TRIGGERS}, "severity":"offen,gering,mittel,hoch", "likelihood":"offen,gering,mittel,hoch", "priority":"offen,1,2,3"}.get(key)
            if choices:
                dv = DataValidation(type="list", formula1='"' + choices + '"', allow_blank=False)
                dv.errorTitle = "Ungültiger Wert"
                dv.error = "Bitte einen Wert der Liste verwenden."
                dv.showErrorMessage = True
                ws.add_data_validation(dv)
                dv.add(f"{get_column_letter(col)}3:{get_column_letter(col)}{max(ws.max_row+100, 103)}")
        ws.row_dimensions[1].hidden = True
    org = wb.create_sheet("Organisation")
    org.append(["Feld", "Wert"])
    for key, value in doc["organization"].items():
        org.append([key, value])
    decisions = wb.create_sheet("Entscheidungen")
    decisions.append(["id", "status", *REVIEW_DEFAULTS])
    for a in doc["activities"]:
        decisions.append([a["id"], review_state(a), *[a["review"][k] for k in REVIEW_DEFAULTS]])
    meta = wb.create_sheet("_meta")
    metadata = {"schema_version": str(SCHEMA), "register_id": doc["register_id"], "revision": str(doc["revision"]), "source_sha256": digest(doc),
                "organization_sha256": digest(doc["organization"]), "reviews_sha256": digest([[a["id"], review_state(a), *[a["review"][k] for k in REVIEW_DEFAULTS]] for a in doc["activities"]])}
    for key, value in metadata.items():
        meta.append([key, value])
    meta.sheet_state = "hidden"
    for ws in wb:
        ws.sheet_view.showGridLines = False
        for row in ws:
            for cell in row:
                if cell.value is not None:
                    value = cell.value
                    if isinstance(value, (date, datetime)):
                        cell.number_format = "yyyy-mm-dd"
                    elif not isinstance(value, int):
                        excel_text(cell, value)
                    cell.font = Font(name="Arial", size=11, color="17324D")
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, ws.max_column+1):
            ws.column_dimensions[get_column_letter(col)].width = 20 if col == 1 else 42
        for row in range(1, ws.max_row+1):
            ws.row_dimensions[row].height = 46 if ws.title != "_meta" else 18
        if ws.title in SHEETS:
            for cell in ws[2]:
                cell.fill = PatternFill("solid", fgColor="234968")
                cell.font = Font(name="Arial", size=11, color="FFFFFF", bold=True)
            for row in ws.iter_rows(min_row=3):
                for cell in row[1:]:
                    cell.fill = PatternFill("solid", fgColor="FFF3CF")
            ws.row_dimensions[2].height = 42
            ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.sheet_properties.outlinePr.summaryRight = False
    overview.column_dimensions["A"].width = 27
    overview.column_dimensions["B"].width = 108
    overview.freeze_panes = "B14"
    overview.row_dimensions[1].height = 28
    # Umgebrochene Eingaben lesbar halten. Sehr lange Zellen bleiben vollständig
    # gespeichert; Excel begrenzt die darstellbare Zeilenhöhe technisch.
    for ws in wb:
        for row in ws:
            lines = 1
            for cell in row:
                if cell.value is not None:
                    width = ws.column_dimensions[cell.column_letter].width or 42
                    lines = max(lines, sum(max(1, math.ceil(len(part) / max(10, width-3))) for part in str(cell.value).split("\n")))
            ws.row_dimensions[row[0].row].height = min(409, max(30, lines*15+10))
    for name in SHEETS:
        wb[name].row_dimensions[1].hidden = True
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    os.chmod(path, 0o600)

def import_xlsx(path, workbook, actor, reason):
    from openpyxl import load_workbook
    source = Path(workbook)
    if source.stat().st_size > MAX_BYTES:
        raise VVTError("Exceldatei überschreitet die technische Grenze.")
    with zipfile.ZipFile(source) as z:
        names = z.namelist()
        if sum(i.file_size for i in z.infolist()) > MAX_BYTES * 8 or any("vbaProject" in n or "externalLink" in n for n in names):
            raise VVTError("Makros, externe Verknüpfungen oder übergroßer Container sind nicht zulässig.")
    wb = load_workbook(source, data_only=False, keep_links=False)
    if set(wb.sheetnames) != {*SHEETS, "Uebersicht", "Organisation", "Entscheidungen", "_meta"}:
        raise VVTError("Erwartete Excelblätter fehlen oder wurden hinzugefügt.")
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == "f" or cell.hyperlink:
                    raise VVTError(f"Formel oder aktiver Link nicht zulässig: {ws.title}!{cell.coordinate}")
    def values(ws):
        return [["" if c.value is None else c.value for c in row] for row in ws]
    meta_rows = values(wb["_meta"])
    if any(len(r) != 2 for r in meta_rows) or len({r[0] for r in meta_rows}) != len(meta_rows):
        raise VVTError("Ungültige Excel-Metadaten.")
    meta = dict(meta_rows)
    if set(meta) != {"schema_version", "register_id", "revision", "source_sha256", "organization_sha256", "reviews_sha256"} or str(meta["schema_version"]) != str(SCHEMA):
        raise VVTError("Unbekannte Excel-Schemaversion.")
    before = read_register(path)
    if meta["register_id"] != before["register_id"]:
        raise VVTError("Exceldatei gehört zu einem anderen Register.")
    org_rows = values(wb["Organisation"])[1:]
    if any(len(r) != 2 for r in org_rows) or len({r[0] for r in org_rows}) != len(org_rows) or digest(dict(org_rows)) != digest(before["organization"]) or meta["organization_sha256"] != digest(before["organization"]):
        raise VVTError("Organisationsblatt wurde verändert. Änderungen gesondert im Registerauftrag behandeln.")
    decision_rows = values(wb["Entscheidungen"])[1:]
    expected_reviews = [[a["id"], review_state(a), *[a["review"][k] for k in REVIEW_DEFAULTS]] for a in before["activities"]]
    if digest(decision_rows) != digest(expected_reviews) or meta["reviews_sha256"] != digest(expected_reviews):
        raise VVTError("Prüfentscheidungen sind nicht über Excel bearbeitbar oder Export ist veraltet.")
    patches = {}
    for name, keys in SHEETS.items():
        rows = values(wb[name])
        if rows[0] != list(keys) or rows[1] != [LABELS.get(k, k) for k in keys]:
            raise VVTError(f"{name}: Spaltenstruktur wurde verändert.")
        ids = set()
        for row in rows[2:]:
            if all(v == "" for v in row):
                continue
            if len(row) != len(keys) or not isinstance(row[0], str) or not row[0] or row[0] in ids:
                raise VVTError(f"{name}: leere, doppelte oder ungültige VT-ID.")
            identifier = row[0]
            ids.add(identifier)
            data = dict(zip(keys, row))
            if name == "Taetigkeiten":
                try:
                    v = data["activity_revision"]
                    if isinstance(v, bool) or str(int(v)) != str(v):
                        raise ValueError()
                    data["activity_revision"] = int(v)
                except (ValueError, TypeError) as exc:
                    raise VVTError("Tätigkeitsrevision muss eine ganze Zahl sein.") from exc
            # Nur das benannte Datumsfeld darf echte Excel-Datumswerte enthalten.
            # Unformatierte Zahlen oder nichtmitternächtliche Uhrzeiten werden nicht geraten.
            if "next_review" in data and isinstance(data["next_review"], (date, datetime)):
                value = data["next_review"]
                if isinstance(value, datetime):
                    if any((value.hour, value.minute, value.second, value.microsecond)) or value.tzinfo is not None:
                        raise VVTError("Nächste Überprüfung darf nur ein Datum ohne Uhrzeit sein.")
                    value = value.date()
                data["next_review"] = value.isoformat()
            if any(not isinstance(v, str) for k, v in data.items() if k != "activity_revision"):
                raise VVTError(f"{name}: Fachfelder müssen Text sein; Prüfdatum als Datum oder YYYY-MM-DD eingeben, keine bloße Seriennummer.")
            patch = patches.setdefault(identifier, {"id": identifier})
            if name == "Screening":
                patch["screening"] = {k:v for k,v in data.items() if k != "id"}
            else:
                patch.update(data)
    known = {a["id"] for a in before["activities"]}
    if any(p["id"] not in known and "title" not in p for p in patches.values()):
        raise VVTError("Neue Tätigkeit muss zuerst im Blatt Taetigkeiten erfasst sein.")
    try:
        revision = int(meta["revision"])
    except (ValueError, TypeError) as exc:
        raise VVTError("Ungültige Exportrevision.") from exc
    return mutate(path, revision, meta["source_sha256"], actor, reason, "import-xlsx", lambda doc: merge_activities(doc, list(patches.values())))

def export_xml(doc, path):
    # JSON-in-XML erhält alle Feldtypen und Leerwerte ohne verlustbehaftete Datumsheuristik.
    root = ET.Element("verarbeitungsverzeichnis", schema_version=str(SCHEMA), register_id=doc["register_id"], revision=str(doc["revision"]), sha256=digest(doc))
    payload = ET.SubElement(root, "register", encoding="application/json")
    payload.text = canonical(doc).decode("utf-8")
    atomic_bytes(path, ET.tostring(root, encoding="utf-8", xml_declaration=True))

def parse_xml(path):
    raw = Path(path).read_bytes()
    if len(raw) > MAX_BYTES:
        raise VVTError("XML-Datei überschreitet die technische Grenze.")
    # Nur UTF-8: dadurch können UTF-16/NUL-Tricks den DTD-Block nicht umgehen.
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise VVTError("XML muss UTF-8-kodiert sein.") from exc
    if "\x00" in text or re.search(r"<!\s*(DOCTYPE|ENTITY)", text, re.I) or re.search(r"^\s*<\?xml[^?]*encoding\s*=\s*['\"](?!utf-8['\"])[^'\"]+", text, re.I):
        raise VVTError("XML mit DTD, Entitäten oder abweichender Kodierung ist nicht zulässig.")
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        raise VVTError("Ungültiges XML.") from exc
    if root.tag != "verarbeitungsverzeichnis" or set(root.attrib) != {"schema_version", "register_id", "revision", "sha256"} or len(root) != 1 or root[0].tag != "register" or root[0].attrib != {"encoding": "application/json"} or len(root[0]):
        raise VVTError("Unbekannte XML-Struktur.")
    doc = validate(json.loads(root[0].text or "", object_pairs_hook=reject_duplicate_keys))
    if root.attrib["schema_version"] != str(SCHEMA) or root.attrib["register_id"] != doc["register_id"] or root.attrib["revision"] != str(doc["revision"]):
        raise VVTError("XML-Metadaten widersprechen dem Registerinhalt.")
    return root.attrib, doc

def import_xml(path, source, actor, reason):
    meta, imported = parse_xml(source)
    before = read_register(path)
    if imported["register_id"] != before["register_id"] or imported["organization"] != before["organization"]:
        raise VVTError("XML gehört zu einer anderen Organisation oder verändert Stammdaten.")
    known = {a["id"]: a for a in before["activities"]}
    for a in imported["activities"]:
        if a["review"] != (known[a["id"]]["review"] if a["id"] in known else REVIEW_DEFAULTS):
            raise VVTError("Prüfentscheidungen werden nicht über XML importiert.")
    patches = [{k:v for k,v in a.items() if k != "review"} for a in imported["activities"]]
    for patch in patches:
        if patch["id"] not in known:
            if patch["activity_revision"] != 1:
                raise VVTError("Neue XML-Tätigkeiten müssen Revision 1 tragen.")
            patch["activity_revision"] = 0
    return mutate(path, int(meta["revision"]), meta["sha256"], actor, reason, "import-xml", lambda doc: merge_activities(doc, patches))

def export_docx(doc, path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml.ns import qn
    output = Document()
    section = output.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = output.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == "Normal" else (15 if name == "Title" else 12))
        style.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rfonts = style.element.get_or_add_rPr().rFonts
        for key in list(rfonts.attrib):
            if "theme" in key.lower():
                del rfonts.attrib[key]
        color = style.element.get_or_add_rPr().find(qn("w:color"))
        if color is not None:
            for key in list(color.attrib):
                if "theme" in key.lower():
                    del color.attrib[key]
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        style.paragraph_format.space_after = Pt(6)
    output.add_heading("Verzeichnis der Verarbeitungstätigkeiten", 0)
    output.add_paragraph(f"{doc['organization']['name']} führt dieses Verzeichnis in Registerrevision {doc['revision']}. Der ausgewiesene Dateistand ist {doc['updated_at']}.")
    output.add_paragraph("Leere Angaben werden ausdrücklich als noch nicht erhoben ausgewiesen. Eine dokumentierte DSFA-Entscheidung ersetzt weder die Datenschutz-Folgenabschätzung selbst noch die rechtliche Zulässigkeitsprüfung der Verarbeitung.")
    output.add_heading("1. Organisation und Kontakt", 1)
    for key, label in (("contact", "Die Organisation ist erreichbar unter"), ("representative", "Als Vertreter ist dokumentiert"), ("dpo", "Zum Datenschutzbeauftragten ist dokumentiert")):
        value = doc["organization"][key] or "noch nicht erhoben"
        output.add_paragraph(f"{label}: {value}.")
    for index, a in enumerate(doc["activities"], 2):
        output.add_page_break()
        output.add_heading(f"{index}. {a['id']} – {a['title']}", 1)
        role = "Verantwortlicher" if a["role"] == "controller" else "Auftragsverarbeiter"
        output.add_paragraph(f"Die dokumentierte Rolle ist {role}. Die Tätigkeit wird mit Status {a['status']} und Tätigkeitsrevision {a['activity_revision']} geführt. Zuständig ist {a['owner'] or 'eine noch zu benennende Person'}.")
        fields = (["purpose", "joint_controller", "data_subjects", "data_categories", "recipients", "transfers", "retention", "toms"] if a["role"] == "controller" else ["controller_contacts", "processing_categories", "transfers", "toms"])
        for n, field in enumerate(fields, 1):
            output.add_heading(f"{index}.{n}. {LABELS[field]}", 2)
            text = a[field].strip()
            output.add_paragraph(f"Dokumentiert ist folgende Angabe: {text}." if text else "Diese Angabe wurde noch nicht erhoben. Die zuständige Person muss den Sachverhalt klären und das Verzeichnis ergänzen.")
        n = len(fields) + 1
        output.add_heading(f"{index}.{n}. Ergänzende Prüfung und Nachweise", 2)
        for field in ("legal_basis", "systems", "processor_contracts", "sources", "next_review"):
            output.add_paragraph(f"Für {LABELS[field]} ist festgehalten: {a[field] or 'noch nicht erhoben'}.")
        output.add_heading(f"{index}.{n+1}. Risiko und Datenschutz-Folgenabschätzung", 2)
        for field, value in a["screening"].items():
            output.add_paragraph(f"Die Angabe zu {LABELS[field]} lautet: {value or 'noch nicht erhoben'}.")
        r = a["review"]
        output.add_paragraph(f"Die gespeicherte Entscheidung lautet {DECISION_LABELS[r['decision']]}; ihr Gültigkeitsstatus ist {review_state(a)}. Eine veraltete Entscheidung darf nicht als aktuelle Freigabe verwendet werden.")
        if r["decision"] != "offen":
            output.add_paragraph(f"{r['reviewer']} hat die Prüfung am {r['reviewed_at']} dokumentiert. Die fallbezogene Begründung lautet: {r['rationale']}. Herangezogen wurden: {r['sources']}.")
        output.add_paragraph("Eine interne Bearbeitungspriorität ist keine gesetzliche Risikoklasse und kann einen gesetzlichen DSFA-Auslöser nicht überstimmen.")
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    output.save(path)
    os.chmod(path, 0o600)

def export_html(doc, path):
    escaped = html.escape
    rows = []
    for f in findings(doc):
        rows.append("<tr>" + "".join("<td>" + escaped(str(v)) + "</td>" for v in [f["id"], f["title"], f["role"], f["review_state"], "; ".join(f["gaps"]) or "Keine im Mindestcheck", f["next_review"]]) + "</tr>")
    content = """<!doctype html><html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>Verarbeitungsverzeichnis</title><style>body{font:16px system-ui;margin:2rem;color:#17324d}table{border-collapse:collapse;width:100%}th,td{text-align:left;padding:.7rem;border-bottom:1px solid #ccd6df}input{padding:.7rem;width:28rem;max-width:90%}small{display:block;margin:1rem 0}a{margin-left:1rem}</style><h1>Verarbeitungsverzeichnis</h1>"""
    content += f"<p>{escaped(doc['organization']['name'])} · Revision {doc['revision']} · {escaped(doc['updated_at'])}</p>"
    content += '<p>Lokale, lesbare Momentaufnahme. Änderungen erfolgen im Register oder über den geprüften Excelimport. Kein Server, keine automatische Cloud-Synchronisation.</p><label>Suche <input id="q" type="search" placeholder="VT-ID, Bezeichnung oder Prüfstatus"></label><button id="download">Register-JSON herunterladen</button><small>Keine Rechtsfreigabe: offene Angaben, veraltete Prüfungen und Pflichttrigger fachlich bearbeiten.</small><table><thead><tr>' + ''.join('<th>'+escaped(v)+'</th>' for v in ['VT-ID','Bezeichnung','Rolle','Prüfung','Lücken','Nächste Prüfung']) + '</tr></thead><tbody>' + ''.join(rows) + '</tbody></table>'
    payload = json.dumps(doc, ensure_ascii=False).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    content += '<script type="application/json" id="data">' + payload + '</script><script>document.getElementById("q").addEventListener("input",e=>{for(const r of document.querySelectorAll("tbody tr"))r.hidden=!r.textContent.toLowerCase().includes(e.target.value.toLowerCase())});document.getElementById("download").addEventListener("click",()=>{const b=new Blob([JSON.stringify(JSON.parse(document.getElementById("data").textContent),null,2)],{type:"application/json"});const u=URL.createObjectURL(b);const a=document.createElement("a");a.href=u;a.download="verarbeitungsverzeichnis.json";a.click();setTimeout(()=>URL.revokeObjectURL(u),1000)});</script></html>'
    atomic_bytes(path, content.encode("utf-8"))

def export(doc, destination, formats):
    validate(doc)
    out = Path(destination)
    out.mkdir(parents=True, exist_ok=True)
    unknown = set(formats) - {"json", "xlsx", "xml", "docx", "html"}
    if unknown:
        raise VVTError("Unbekanntes Exportformat.")
    paths = [out / ("verarbeitungsverzeichnis." + f) for f in formats]
    if any(p.exists() for p in paths):
        raise VVTError("Exportziel enthält bereits eine Ausgabedatei. Neues Verzeichnis verwenden.")
    for f, p in zip(formats, paths):
        if f == "json":
            atomic_bytes(p, canonical(doc))
        else:
            globals()["export_" + f](doc, p)
    return [str(p) for p in paths]

def main(argv=None):
    parser = argparse.ArgumentParser(description="Lokales Verarbeitungsverzeichnis; keine automatische Rechtsfreigabe oder Cloud-Synchronisation.")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("init", help="Leeres Register anlegen")
    p.add_argument("register"); p.add_argument("--organization", required=True)
    for key in ("contact", "representative", "dpo"):
        p.add_argument("--"+key, default="")
    p.add_argument("--actor", required=True); p.add_argument("--reason", required=True)
    for command in ("validate", "status"):
        p = sub.add_parser(command); p.add_argument("register")
    p = sub.add_parser("export"); p.add_argument("register"); p.add_argument("--out", required=True)
    p.add_argument("--formats", default="json,xlsx,xml,docx,html")
    for command in ("upsert", "review", "update-org", "import-xlsx", "import-xml"):
        p = sub.add_parser(command); p.add_argument("register"); p.add_argument("--actor", required=True); p.add_argument("--reason", required=True)
        if command.startswith("import-"):
            p.add_argument("--input", required=True)
        else:
            p.add_argument("--expected-revision", type=int, required=True); p.add_argument("--expected-sha256", required=True)
        if command in ("upsert", "update-org"):
            p.add_argument("--input", required=True, help="JSON-Objekt oder Liste mit VT-ID und geänderten Sachfeldern")
        if command == "review":
            p.add_argument("--id", required=True); p.add_argument("--decision", choices=("erforderlich", "begruendet_nicht_erforderlich"), required=True)
            p.add_argument("--rationale", required=True); p.add_argument("--sources", required=True)
    args = parser.parse_args(argv)
    try:
        if args.command == "init":
            doc = create(args.register, {"name":args.organization, "contact":args.contact, "representative":args.representative, "dpo":args.dpo}, args.actor, args.reason)
        elif args.command == "upsert":
            data = read_json(args.input)
            doc = upsert(args.register, data if isinstance(data, list) else [data], args.expected_revision, args.expected_sha256, args.actor, args.reason)
        elif args.command == "update-org":
            doc = update_organization(args.register, read_json(args.input), args.expected_revision, args.expected_sha256, args.actor, args.reason)
        elif args.command == "review":
            doc = review(args.register, args.id, args.decision, args.rationale, args.sources, args.expected_revision, args.expected_sha256, args.actor, args.reason)
        elif args.command == "import-xlsx":
            doc = import_xlsx(args.register, args.input, args.actor, args.reason)
        elif args.command == "import-xml":
            doc = import_xml(args.register, args.input, args.actor, args.reason)
        else:
            doc = read_register(args.register)
        if args.command == "export":
            output = {"files": export(doc, args.out, args.formats.split(",")), "revision":doc["revision"], "sha256":digest(doc)}
        else:
            output = {"register_id":doc["register_id"], "revision":doc["revision"], "sha256":digest(doc), "activities":len(doc["activities"]), "findings":findings(doc)}
        print(json.dumps(output, ensure_ascii=False, indent=2))
        return 0
    except (VVTError, OSError, ValueError, zipfile.BadZipFile, ImportError) as exc:
        print("VVT: " + str(exc), file=sys.stderr)
        return 2

if __name__ == "__main__":
    sys.exit(main())
