#!/usr/bin/env python3
"""Baut ausschließlich diese Zusatzakte; keine Marketplace- oder Indexänderung."""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import io
import json
import re
import sys
import tempfile
import zipfile
from xml.sax.saxutils import escape
from datetime import datetime, timezone
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from akteninhalt import CSVS, DOCS, MAILS, TEXTS

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
sys.path.insert(0, str(REPO / "scripts"))
from testakte_disclaimer import NOTICE_DE, NOTICE_EN, pdf_content_errors
from testakte_office_pdf import render_office_batch

ORIGINALS = HERE / "akte" / "originalunterlagen"
DOWNLOADS = HERE / "downloads"
NOTICE = f"{NOTICE_DE}\n\n{NOTICE_EN}\n\nGVB-REI-2026-017, Los 1. Aktenstand: 02.10.2026.\n"


def load_pdf_builder():
    spec = importlib.util.spec_from_file_location("einzelpdf", REPO / "scripts/build-testakten-einzelpdf-zips.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def word_document(item):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.3)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    for name, size in (("Title", 15), ("Heading 1", 12)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        properties = style.element.get_or_add_pPr()
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
        borders = OxmlElement("w:pBdr")
        for edge in ("top", "left", "bottom", "right", "between"):
            border = OxmlElement("w:" + edge)
            border.set(qn("w:val"), "nil")
            borders.append(border)
        properties.append(borders)
    for index, line in enumerate(item["sender"]):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(11 if index == 0 else 9)
        run.bold = index == 0
    p = doc.add_paragraph(item["recipient"])
    p.paragraph_format.space_before = Pt(15)
    doc.add_paragraph(f"Berlin, {item['date']} | GVB-REI-2026-017")
    doc.add_paragraph(item["title"], style="Title")
    new_page = {"05": "4. Kontrolle und Leistungsnachweis", "06": "3. Ausführungskonzept",
                "07": "4. Entgeltanpassung"}.get(item["name"][:2])
    for text in item["paragraphs"]:
        if re.fullmatch(r"\d+\. [^\n]+", text):
            p = doc.add_paragraph(text, style="Heading 1")
            if text == new_page:
                p.paragraph_format.page_break_before = True
        else:
            p = doc.add_paragraph(text)
            p.paragraph_format.widow_control = True
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph(item["signature"])
    p.paragraph_format.keep_together = True
    footer = section.footer.paragraphs[0]
    footer.add_run("GVB-REI-2026-017 | " + item["name"].split("_", 1)[0] + " | Seite ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.name, run.font.size = "Arial", Pt(8)
    core = doc.core_properties
    core.author = "Klotzkette"
    core.title = item["title"]
    core.created = core.modified = datetime(2026, 10, 2, tzinfo=timezone.utc)
    path = ORIGINALS / f"{item['name']}.docx"
    doc.save(path)
    # Office ZIP timestamps otherwise vary even when the visible content is identical.
    with zipfile.ZipFile(path) as archive:
        members = [(name, archive.read(name)) for name in archive.namelist()]
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in members:
            write_member(archive, name, data)


def write_member(archive, name, data):
    info = zipfile.ZipInfo(name, (2026, 10, 2, 0, 0, 0))
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o100644 << 16
    archive.writestr(info, data)


def csv_pdf(source):
    labels = {
        "02": ("Objektliste vom 17. Juli 2026", [140, 150, 35, 55, 80, 85, 75]),
        "03": ("Kostenansatz Finanzen vom 20. Juli 2026", [35, 180, 95, 45, 110, 155]),
        "22": ("Abschließende Preisangaben vom 22. September 2026", [155, 105, 105, 45, 115, 95]),
        "34": ("Portalprotokoll bis 1. Oktober 2026", [80, 65, 250, 170, 155]),
        "38": ("Leistungsverzeichnis und unbepreistes Preisblatt", [60, 215, 110, 95, 100, 140]),
    }
    title, widths = labels[source.name[:2]]
    rows = list(csv.reader(io.StringIO(source.read_text(encoding="utf-8-sig")), delimiter=";"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=10, leading=13)
    header = ParagraphStyle("header", parent=cell, fontName="Helvetica-Bold")
    rendered = []
    for index, row in enumerate(rows):
        rendered.append([Paragraph(escape(value.replace("_", " ").replace("Durchgänge/Jahr", "Durchgänge pro Jahr") if index == 0 else value), header if index == 0 else cell)
                         for value in row])
    table = Table(rendered, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c7ced1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5edf0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9fa")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
    ]))
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=42, rightMargin=42,
                            topMargin=38, bottomMargin=40, title=title, author="Klotzkette", invariant=1)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(42, 24, "GVB-REI-2026-017 | " + source.name)
        canvas.drawRightString(800, 24, f"Seite {document.page}")
        canvas.restoreState()

    doc.build([Paragraph("GVB Gemeinsame Verkehrsbetriebe Berlin AöR", header), Spacer(1, 9),
               Paragraph(title, ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=16, leading=20)),
               Spacer(1, 15), table], onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def build_originals():
    ORIGINALS.mkdir(parents=True, exist_ok=True)
    for item in DOCS:
        word_document(item)
    for number, slug, date, sender, recipient, subject, body in MAILS:
        mail = EmailMessage(policy=SMTP)
        mail["From"], mail["To"], mail["Subject"] = sender, recipient, subject
        mail["Date"] = datetime.fromisoformat(date)
        mail["Message-ID"] = f"<gvb-017-{number:02d}@poststelle.gvb-berlin.example>"
        mail.set_content(body + "\n", charset="utf-8")
        (ORIGINALS / f"{number:02d}_{slug}.eml").write_bytes(mail.as_bytes())
    for name, body in TEXTS.items():
        (ORIGINALS / name).write_text(body, encoding="utf-8")
    for name, rows in CSVS.items():
        with (ORIGINALS / name).open("w", newline="", encoding="utf-8-sig") as stream:
            csv.writer(stream, delimiter=";").writerows(rows)


def build():
    build_originals()
    sources = sorted(ORIGINALS.iterdir())
    expected = len(DOCS) + len(MAILS) + len(TEXTS) + len(CSVS)
    if len(sources) != expected:
        raise ValueError("Unerwartete Datei im Originalordner; keine stillschweigende Auslieferung")
    DOWNLOADS.mkdir(exist_ok=True)
    renderer = load_pdf_builder()
    original_footer = renderer.G.header_footer_factory
    renderer.G.header_footer_factory = lambda _: original_footer("GVB-REI-2026-017")
    office_cache = render_office_batch(sources)
    missing = [p.name for p in sources if p.suffix == ".docx" and p not in office_cache]
    if missing:
        raise RuntimeError(f"Native Word-Konvertierung fehlt: {missing}")
    writer = PdfWriter()
    records = []
    # Publish the three formats only after every source has rendered successfully.
    with tempfile.TemporaryDirectory(prefix="gvb-ausgabe-") as temporary:
        staging = Path(temporary)
        with zipfile.ZipFile(staging / "GVB_Reinigungsvergabe_Einzel_PDFs.zip", "w") as pdf_zip, \
             zipfile.ZipFile(staging / "GVB_Reinigungsvergabe_Originale.zip", "w") as original_zip:
            for archive in (pdf_zip, original_zip):
                write_member(archive, "README.txt", NOTICE.encode("utf-8"))
            for source in sources:
                data = csv_pdf(source) if source.suffix == ".csv" else renderer.render_document_pdf(source, ORIGINALS, office_cache)
                if not data or pdf_content_errors(data):
                    raise ValueError(f"Ungültiges Dokument: {source.name}")
                reader = PdfReader(io.BytesIO(data))
                pdf_name = source.stem + ".pdf"
                first_page = len(writer.pages) + 1
                writer.append(reader, outline_item=source.name)
                write_member(pdf_zip, pdf_name, data)
                write_member(original_zip, source.name, source.read_bytes())
                records.append({"datei": source.name, "pdf": pdf_name, "seiten": len(reader.pages),
                                "gesamt_ab_seite": first_page,
                                "sha256": hashlib.sha256(source.read_bytes()).hexdigest()})
        writer.add_metadata({"/Title": "GVB Reinigungsvergabe 2026 | Los 1", "/Author": "Klotzkette"})
        writer.write(staging / "GVB_Reinigungsvergabe_Gesamt.pdf")
        for file in staging.iterdir():
            (DOWNLOADS / file.name).write_bytes(file.read_bytes())
    (HERE / "aktenbestand.json").write_text(json.dumps({"aktenstand": "2026-10-02", "dokumente": records},
                                                      ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{expected} Originalunterlagen; {len(writer.pages)} Seiten; drei übereinstimmende Ausgabeformate")


if __name__ == "__main__":
    build()
