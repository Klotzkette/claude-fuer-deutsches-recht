#!/usr/bin/env python3
"""Prüft diese Zusatzsammlung ohne Änderungen an zentralen Verzeichnissen."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "scripts"))
from testakte_disclaimer import NOTICE_DE, NOTICE_EN, pdf_content_errors


def require(condition, message):
    if not condition:
        raise ValueError(message)


def normalized(text):
    return re.sub(r"\s+", " ", text).strip()


def check():
    original_dir = HERE / "akte/originalunterlagen"
    files = sorted(original_dir.iterdir())
    inventory = json.loads((HERE / "aktenbestand.json").read_text())
    entries = inventory["dokumente"]
    require(len(files) == len(entries) == 38, "Aktenbestand muss 38 Dokumente umfassen")
    require([f.name for f in files] == [r["datei"] for r in entries], "Inventar weicht ab")
    require({int(f.name[:2]) for f in files} == set(range(1, 39)), "Dokumentnummer fehlt")
    counts = {ext: sum(p.suffix == ext for p in files) for ext in (".docx", ".eml", ".csv", ".txt")}
    require(counts == {".docx": 20, ".eml": 10, ".csv": 5, ".txt": 3}, "Formatbestand weicht ab")

    merged = PdfReader(HERE / "downloads/GVB_Reinigungsvergabe_Gesamt.pdf")
    require(len(merged.pages) == sum(r["seiten"] for r in entries), "Seitenzahl weicht ab")
    cursor = 0
    with zipfile.ZipFile(HERE / "downloads/GVB_Reinigungsvergabe_Originale.zip") as native, \
         zipfile.ZipFile(HERE / "downloads/GVB_Reinigungsvergabe_Einzel_PDFs.zip") as pdfs:
        for archive in (native, pdfs):
            require(archive.testzip() is None, "Defektes ZIP")
            names = archive.namelist()
            require(len(names) == len(set(names)) == 39, "ZIP enthält Dubletten oder falsche Anzahl")
            require(all("/" not in n and "\\" not in n for n in names), "ZIP ist nicht flach")
            notice = archive.read("README.txt").decode("utf-8")
            require(notice.startswith(NOTICE_DE + "\n\n" + NOTICE_EN), "ZIP-Hinweis fehlt am Anfang")
            require(not any(n.lower().endswith(".md") for n in names), "Markdown im Akten-ZIP")
        require(set(native.namelist()) == {p.name for p in files} | {"README.txt"}, "Originale fehlen")
        require(set(pdfs.namelist()) == {r["pdf"] for r in entries} | {"README.txt"}, "Einzel-PDF fehlt")
        for file, record in zip(files, entries):
            data = file.read_bytes()
            require(native.read(file.name) == data, f"Veraltetes Original: {file.name}")
            require(hashlib.sha256(data).hexdigest() == record["sha256"], "Falsche Prüfsumme")
            pdf = pdfs.read(record["pdf"])
            require(not pdf_content_errors(pdf), f"Unzulässiger PDF-Hinweis: {file.name}")
            reader = PdfReader(io.BytesIO(pdf))
            require(len(reader.pages) == record["seiten"], f"PDF-Seiten fehlen: {file.name}")
            require(record["gesamt_ab_seite"] == cursor + 1, "Seitenzuordnung fehlerhaft")
            for page in reader.pages:
                text = normalized(page.extract_text() or "")
                require(len(text) > 80, f"Leere oder fast leere PDF-Seite: {file.name}")
                require(text == normalized(merged.pages[cursor].extract_text() or ""),
                        f"Gesamt-PDF weicht von Einzel-PDF ab: {file.name}")
                cursor += 1
            pdf_text = " ".join(page.extract_text() or "" for page in reader.pages)
            pdf_text = normalized(re.sub(r"GVB-REI-2026-017\s*\|\s*\d+\s*\|\s*Seite\s+\d+", "", pdf_text))
            if file.suffix == ".docx":
                doc = Document(file)
                text = "\n".join(p.text for p in doc.paragraphs)
                require(len(text) > 1000, f"Unvollständiges Word-Dokument: {file.name}")
                for paragraph in doc.paragraphs:
                    for line in paragraph.text.splitlines():
                        if line.strip():
                            require("".join(line.split()) in "".join(pdf_text.split()),
                                    f"Textverlust im PDF: {file.name}: {line[:60]}")
            elif file.suffix == ".eml":
                mail = BytesParser(policy=policy.default).parsebytes(data)
                require(not mail.defects, f"Defekte E-Mail: {file.name}")
                for header in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version", "Content-Type"):
                    require(bool(mail[header]), f"E-Mail-Header fehlt: {file.name}/{header}")
                text = mail.get_body(preferencelist=("plain",)).get_content()
                require(len(text) > 450, f"E-Mail zu knapp: {file.name}")
            elif file.suffix == ".csv":
                text = data.decode("utf-8-sig")
                rows = list(csv.reader(io.StringIO(text), delimiter=";"))
                require(len(rows) >= 5, "Zu wenig Datenzeilen")
                require(all(len(row) == len(rows[0]) for row in rows), "CSV-Spalten verrutscht")
            else:
                text = data.decode("utf-8")
            require(not re.search(r"Testakte|fiktiv|Platzhalter|Formathinweis|Lösungsmatrix|§", text, re.I),
                    f"Unzulässiger Aktenhinweis: {file.name}")

    rows = list(csv.DictReader((original_dir / "03_Kostenansatz_Finanzen.csv").open(encoding="utf-8-sig"), delimiter=";"))
    require(sum(Decimal(r["Gesamt_EUR_netto"]) for r in rows) == 7000000, "Kostenansatz nicht 7 Mio.")
    require(sum(Decimal(r["Gesamt_EUR_netto"]) for r in rows if r["Los"] == "1") == 4020000,
            "Los-1-Schätzwert weicht von Bekanntmachung ab")
    for row in rows:
        require(Decimal(row["Jahr_EUR_netto"]) * Decimal(row["Jahre"]) == Decimal(row["Gesamt_EUR_netto"]),
                "Jahres- und Gesamtwert widersprechen sich")
    prices = list(csv.DictReader((original_dir / "22_Abschliessende_Preisangaben.csv").open(encoding="utf-8-sig"), delimiter=";"))
    totals = {}
    for row in prices:
        require(int(row["Jahr_EUR_netto"]) * int(row["Jahre"]) == int(row["Gesamt_EUR_netto"]), "Preisblattfehler")
        totals[row["Bieter"]] = totals.get(row["Bieter"], 0) + int(row["Gesamt_EUR_netto"])
    require(totals == {"Spreeklar": 3880000, "Märkischer Objektservice": 4010000, "Nordlicht": 4165000},
            "Schlussangebote widersprechen dem Preisblatt")
    require(26400 * Decimal("16.40") + sum((58000, 100000, 54000, 38000, 22000, 30000, 13040)) == 748000,
            "Preisaufklärung nicht geschlossen")
    objects = list(csv.DictReader((original_dir / "02_Objektliste_17_07.csv").open(encoding="utf-8-sig"), delimiter=";"))
    quantities = list(csv.DictReader((original_dir / "38_Leistungsverzeichnis_Preisblatt.csv").open(encoding="utf-8-sig"), delimiter=";"))
    for obj, quantity in zip(objects[:5], quantities[:5]):
        require(int(obj["Menge"]) * int(obj["Durchgänge/Jahr"]) == int(quantity["Menge_Jahr"]),
                "Leistungsverzeichnis widerspricht der Objektliste")

    workflows = sorted((HERE / "workflows").glob("*.md"))
    require(len(workflows) == 10, "Es müssen zehn Workflows sein")
    lengths = []
    for workflow in workflows:
        text = workflow.read_text(encoding="utf-8")
        require(max(len(text), len(text.replace("\n", "\r\n")), len(text.encode("utf-8")),
                    len(text.encode("utf-16-le")) // 2) <= 7500, f"Längengrenze: {workflow.name}")
        for field in ("SCHRITT:", "STATUS:", "ERGEBNISSE:", "OFFEN:", "FRISTEN:", "WEITER:", "FREIGABE:"):
            require(field in text, f"Übergabefeld fehlt: {workflow.name}/{field}")
        require("RUECKFRAGE" in text and "https://" in text, "Rückfrage oder Quellen fehlen")
        lengths.append(len(text))
    require(not list(HERE.rglob("SKILL.md")), "Zusatzsammlung darf kein installierbarer Skill sein")
    require(not list(HERE.rglob("plugin.json")), "Zusatzsammlung darf kein Plugin sein")
    readme = (HERE / "README.md").read_text()
    require(readme.index(NOTICE_DE) < readme.index("downloads/"), "Hinweis steht nicht vor Downloads")
    for target in re.findall(r"\]\(([^)]+)\)", readme):
        if not target.startswith("https://"):
            require((HERE / target.split("?", 1)[0]).exists(), f"Toter Link: {target}")
    print(f"OK: 38 Originale, 38 Einzel-PDFs, {cursor} Gesamtseiten; flache ZIPs; Preise; Quellenzuordnung.")
    print(f"OK: zehn Workflows, {min(lengths)} bis {max(lengths)} Zeichen; keine Plugin-Registrierung.")


if __name__ == "__main__":
    check()
