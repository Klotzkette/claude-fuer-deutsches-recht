#!/usr/bin/env python3
"""Prüft Umfang, Hauptworkflows und die sechs nativen Wirtschaftsmandate."""

from __future__ import annotations

import csv
from email import policy
from email.parser import BytesParser
import importlib.util
from pathlib import Path
import re
import unittest

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageStat
from pypdf import PdfReader
import yaml

from quality_lab import load, marketplace, validate_profile
from testakte_office_pdf import uncached_formula_cells
from testakte_zip_common import working_dump_flat_pairs


ROOT = Path(__file__).resolve().parent.parent
PACKAGES = {
    "anwaltschaft-generell": (10, "mandat-bis-zum-schreiben-bearbeiten", (
        "anwaltschaft-lieferstreit-kaffeeroesterei-leipzig",
        "anwaltschaft-arbeitsrecht-vertrieb-mainz",
        "anwaltschaft-dienstleister-datenzugang-dortmund",
    )),
    "corporate-contract-law": (20, "wirtschaftsvertrag-bis-zur-unterschrift-erstellen", (
        "corporate-contract-law-rahmenlieferung-sensorik-aachen",
        "corporate-contract-law-projektvertrag-automation-augsburg",
        "corporate-contract-law-vertrieb-messtechnik-bremen",
    )),
}


def source_text(path):
    if path.suffix == ".docx":
        document = Document(path)
        return "\n".join([p.text for p in document.paragraphs] + [
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        ])
    if path.suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        body = message.get_body(preferencelist=("plain",))
        return body.get_content() if body else ""
    if path.suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if path.suffix in {".csv", ".txt"}:
        return path.read_text(encoding="utf-8-sig")
    return ""


class PracticePackageTests(unittest.TestCase):
    def test_exact_skill_counts_and_main_skill_remain_installable(self):
        entries = marketplace()
        for name, (count, main, _) in PACKAGES.items():
            with self.subTest(plugin=name):
                directory = entries[name]
                paths = list((directory / "skills").glob("*/SKILL.md"))
                self.assertEqual(len(paths), count)
                self.assertTrue((directory / "skills" / main / "SKILL.md").is_file())
                profile = load(ROOT / "quality/evals" / f"{name}.json")
                validate_profile(profile, name, directory)
                self.assertEqual(profile["selection"]["target_skill"], main)
                self.assertEqual(profile["cases"][0]["target_skill"], main)
                self.assertEqual({case["target_skill"] for case in profile["cases"]},
                                 {path.parent.name for path in paths})
                for path in paths:
                    front = yaml.safe_load(path.read_text().split("---", 2)[1])
                    self.assertEqual(front["name"], path.parent.name)
                    self.assertEqual(set(front), {"name", "description"})
                    self.assertIn("zitierweise.md", path.read_text())
                    self.assertRegex(path.read_text(), r"(?m)^# 1\. [^\n]+$")

    def test_legal_review_boundaries_remain_explicit(self):
        general = (ROOT / "anwaltschaft-generell/anwaltschaft-generell-werkstatt.md").read_text()
        corporate = (ROOT / "corporate-contract-law/corporate-contract-law-werkstatt.md").read_text()
        self.assertIn("Ein Zeuge vom Hörensagen ist nicht allein deshalb unzulässig", general)
        self.assertNotIn("Ein Zeuge muss den behaupteten Vorgang selbst wahrgenommen haben", general)
        self.assertIn("Paragraf 92c Absatz 1 HGB", corporate)
        self.assertIn("des Unionsrechts oder des Rechts eines Mitgliedstaats", corporate)
        self.assertNotIn("soweit er nicht durch anwendbares Recht", corporate)

    def test_focus_discovery_is_not_limited_to_specialist_titles(self):
        spec = importlib.util.spec_from_file_location(
            "focus_coverage", ROOT / "scripts/test-schwerpunkt-coverage.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        module.SchwerpunktCoverage.setUpClass()
        self.assertTrue(set(PACKAGES) <= module.SchwerpunktCoverage.packages.keys())
        overview = (ROOT / "SCHWERPUNKTE.md").read_text()
        for name in PACKAGES:
            self.assertIn(f"{name}-hauptproblem.md", overview)

    def test_three_standalone_prompts_are_separate_and_substantial(self):
        for name in PACKAGES:
            directory = ROOT / name
            for kind in ("werkstatt", "schnellstart", "hauptproblem"):
                path = directory / f"{name}-{kind}.md"
                text = path.read_text()
                size = len(text.encode("utf-8"))
                with self.subTest(prompt=path.name):
                    self.assertGreater(size, 3000)
                    if kind == "werkstatt":
                        self.assertGreater(size, 35000)
                        self.assertLessEqual(size, 128 * 1024)
                    else:
                        self.assertLessEqual(size, 7500)
                    self.assertIn("BGB", text)
                    self.assertRegex(text, r"BGH|BAG|EuGH")
                    self.assertNotIn(chr(167), text)
                    self.assertTrue(all(re.match(r"\d+(?:\.\d+)*\.?\s", h)
                                        for h in re.findall(r"^#{1,6}\s+(.+)$", text, re.M)))
                    self.assertFalse(list((directory / "skills").rglob(path.name)))

    def test_each_case_has_native_documents_not_only_export_copies(self):
        for name, (_, _, slugs) in PACKAGES.items():
            for slug in slugs:
                directory = ROOT / "testakten" / slug
                pairs = working_dump_flat_pairs(directory, include_gesamt_pdf=False)
                with self.subTest(case=slug):
                    self.assertGreaterEqual(len(pairs), 20)
                    self.assertTrue({".docx", ".eml", ".xlsx", ".csv", ".png"}
                                    <= {path.suffix for path, _ in pairs})
                    self.assertFalse(any("/" in arc or "\\" in arc for _, arc in pairs))
                    self.assertFalse(any(path.suffix in {".md", ".yaml"} for path, _ in pairs))
                    self.assertEqual(len({arc.casefold() for _, arc in pairs}), len(pairs))
                    self.assertIn(f"`{name}`", (directory / "README.md").read_text())
                    self.assertTrue((directory / "rubric.yaml").is_file())
                for path, _ in pairs:
                    text = source_text(path)
                    for marker in ("Lösungsmatrix", "Musterlösung", "Testakte", "Formathinweis", "[TODO]", chr(167)):
                        self.assertNotIn(marker, text, path)
                    if path.suffix in {".docx", ".eml", ".pdf"}:
                        self.assertGreaterEqual(len(text), 600, path)

    def test_mails_have_complete_headers_and_csvs_have_stable_columns(self):
        for _, _, slugs in PACKAGES.values():
            for slug in slugs:
                directory = ROOT / "testakten" / slug
                for path in directory.glob("*.eml"):
                    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                    for header in ("From", "To", "Date", "Subject", "Message-ID", "Content-Type"):
                        self.assertTrue(message.get(header), (path, header))
                    self.assertNotIn("\ufffd", source_text(path), path)
                for path in directory.glob("*.csv"):
                    with path.open(encoding="utf-8-sig", newline="") as stream:
                        rows = list(csv.reader(stream, delimiter=";", strict=True))
                    self.assertGreaterEqual(len(rows), 5, path)
                    self.assertTrue(all(len(row) == len(rows[0]) for row in rows), path)

    def test_workbooks_have_recalculated_formulas_and_visuals_are_nonblank(self):
        for _, _, slugs in PACKAGES.values():
            for slug in slugs:
                directory = ROOT / "testakten" / slug
                for path in directory.glob("*.xlsx"):
                    self.assertEqual(uncached_formula_cells(path), [], path)
                    cached = load_workbook(path, data_only=True)
                    formulas = load_workbook(path, data_only=False)
                    count = 0
                    try:
                        for sheet in formulas:
                            self.assertTrue(sheet.freeze_panes, (path, sheet.title))
                            for row in sheet:
                                for cell in row:
                                    if cell.data_type == "f":
                                        count += 1
                                        value = cached[sheet.title][cell.coordinate]
                                        self.assertNotEqual(value.data_type, "e", (path, cell.coordinate))
                        self.assertGreater(count, 0, path)
                    finally:
                        cached.close()
                        formulas.close()
                for path in directory.glob("*.png"):
                    with Image.open(path) as picture:
                        self.assertGreaterEqual(picture.width, 900, path)
                        self.assertGreater(ImageStat.Stat(picture.convert("L")).stddev[0], 15, path)

    def test_case_amounts_and_alternative_party_positions_remain_distinct(self):
        checks = {
            "anwaltschaft-lieferstreit-kaffeeroesterei-leipzig/25_zahlungen_px12.xlsx": {
                "Zahlungen": {"E6": 42840, "E7": 14280, "E8": 28560, "E15": 28560},
            },
            "anwaltschaft-arbeitsrecht-vertrieb-mainz/15_bonusabgleich_2025.xlsx": {
                "Bonus 2025": {"B27": 710000, "C27": 650000, "E27": 14200,
                               "F27": 13000, "F31": 1800, "F32": 11200},
            },
            "anwaltschaft-dienstleister-datenzugang-dortmund/23_rechnungs_leistungsabgleich.xlsx": {
                "Rechnungsabgleich": {"H5": 3272.50, "D18": 2750, "G13": 10710},
                "Stunden": {"D10": 12, "F10": 1500, "D13": 4, "F13": 500},
            },
            "corporate-contract-law-rahmenlieferung-sensorik-aachen/12_preisstaffeln_d400.xlsx": {
                "Preisstaffeln": {"D11": 180.42, "E11": 216504, "B18": 3600},
                "Anlauf": {"B13": 240, "G15": 14756, "G18": 14756, "G19": 54454.40},
            },
            "corporate-contract-law-projektvertrag-automation-augsburg/17_Meilensteinzahlungen_und_Termine.xlsx": {
                "Zahlungen": {"B9": 11305, "C17": 675000, "E17": 803250,
                              "C23": 125500, "C25": 135000, "C26": 684500},
                "Zeitplan": {"D8": 28},
            },
            "corporate-contract-law-vertrieb-messtechnik-bremen/13_Haendlerkalkulation_2027.xlsx": {
                "Kalkulation": {"E15": 35550, "E23": 180000, "E24": 150000,
                                "E30": 10665, "E31": 14220, "E32": 10665},
                "Monatsforecast": {"B18": 340, "C18": 87, "F18": 202830},
            },
        }
        for relative, sheets in checks.items():
            path = ROOT / "testakten" / relative
            workbook = load_workbook(path, data_only=True)
            try:
                for title, cells in sheets.items():
                    for coordinate, expected in cells.items():
                        with self.subTest(file=relative, sheet=title, cell=coordinate):
                            self.assertAlmostEqual(workbook[title][coordinate].value,
                                                   expected, places=2)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
