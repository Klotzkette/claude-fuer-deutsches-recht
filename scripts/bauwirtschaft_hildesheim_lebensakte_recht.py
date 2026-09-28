"""Nicht vollzogene Bauträger-Verkaufsoption zum Hildesheimer Mietprojekt.

Autor: Klotzkette. Quellenprüfung: quality/hildesheim-lebensakte/rechtsprechung-bautraeger.md.
Die echte Fallchronologie und ihre 198 Originale werden weder umgeschrieben noch bebucht.
"""
from __future__ import annotations

import csv
import io
import json
from decimal import Decimal
from pathlib import Path

from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from bauwirtschaft_hildesheim_lebensakte_common import (
    CASE, QUALITY, REF, euro, new_document, add_table, save_docx, written_norms,
)

OPTION = Path('13_Option_Bautraegerverkauf')
STAND = '2027-11-01'
GRUNDSTUECK = ('Grundbuch von Hildesheim Blatt 88018, Gemarkung Hildesheim, '
               'Flur 88, Flurstück 18/6, Gebäude- und Freifläche Wohnhof Am Steinbogen 18, '
               'Größe 1.300 m²')
UNITS = [
    dict(n=1, area=68, floor='Erdgeschoss', side='westlicher Gebäudeteil', price=409000,
         names='Johann Müller und Anna Müller', address='Lindensteg 12, 31134 Hildesheim',
         born='Johann Müller, geboren am 12.03.1968, und Anna Müller, geboren am 21.06.1970',
         purpose='gemeinschaftliche Eigennutzung', married=True),
    dict(n=2, area=75, floor='Erdgeschoss', side='mittlerer Gebäudeteil', price=449000,
         names='Georg Schmidt und Margarete Schmidt', address='Wiesenbogen 7, 31135 Hildesheim',
         born='Georg Schmidt, geboren am 07.02.1961, und Margarete Schmidt, geboren am 19.09.1963',
         purpose='private Vermietung nach Übergabe', married=True),
    dict(n=3, area=82, floor='Erdgeschoss', side='östlicher Gebäudeteil', price=489000,
         names='Friedrich Schulz und Erna Schulz', address='Buchensteg 19, 31137 Hildesheim',
         born='Friedrich Schulz, geboren am 08.11.1972, und Erna Schulz, geboren am 17.04.1974',
         purpose='gemeinschaftliche Eigennutzung', married=True),
    dict(n=4, area=68, floor='erstes Obergeschoss', side='westlicher Gebäudeteil', price=419000,
         names='Hans Becker', address='Ulmenbogen 22, 31134 Hildesheim',
         born='Hans Becker, geboren am 24.01.1965', purpose='Eigennutzung', married=False),
    dict(n=5, area=75, floor='erstes Obergeschoss', side='mittlerer Gebäudeteil', price=459000,
         names='Karl Hoffmann und Hildegard Hoffmann', address='Eschensteg 8, 31135 Hildesheim',
         born='Karl Hoffmann, geboren am 03.05.1960, und Hildegard Hoffmann, geboren am 26.08.1962',
         purpose='private Vermietung nach Übergabe', married=True),
    dict(n=6, area=82, floor='erstes Obergeschoss', side='östlicher Gebäudeteil', price=499000,
         names='Heinrich Koch und Elisabeth Koch', address='Weidenbogen 5, 31137 Hildesheim',
         born='Heinrich Koch, geboren am 16.07.1971, und Elisabeth Koch, geboren am 14.12.1973',
         purpose='gemeinschaftliche Eigennutzung', married=True),
    dict(n=7, area=75, floor='zweites Obergeschoss', side='westlicher Gebäudeteil', price=469000,
         names='Otto Fischer und Martha Fischer', address='Eichensteg 16, 31134 Hildesheim',
         born='Otto Fischer, geboren am 27.10.1959, und Martha Fischer, geboren am 06.03.1961',
         purpose='gemeinschaftliche Eigennutzung', married=True),
    dict(n=8, area=75, floor='zweites Obergeschoss', side='östlicher Gebäudeteil', price=479000,
         names='Wilhelm Weber und Gertrud Weber', address='Ahornbogen 11, 31135 Hildesheim',
         born='Wilhelm Weber, geboren am 11.06.1966, und Gertrud Weber, geboren am 23.02.1967',
         purpose='private Vermietung nach Übergabe', married=True),
]
RATES = [
    ('1', 'Beginn der Erdarbeiten', Decimal('30'), 'Grundstücksanteil und Beginn der Erdarbeiten.'),
    ('2', 'Fertigstellung des Rohbaus einschließlich Zimmererarbeiten', Decimal('28'), '40 Prozent der verbleibenden 70 Prozent.'),
    ('3', 'Herstellung der Dachflächen und Dachrinnen sowie Fenstereinbau einschließlich Verglasung', Decimal('12.6'), '8 und 10 Prozent der verbleibenden 70 Prozent; beide Leistungen müssen fertiggestellt sein.'),
    ('4', 'Vollständige Rohinstallation von Heizung, Sanitär und Elektro', Decimal('6.3'), 'Je 3 Prozent der verbleibenden 70 Prozent; alle drei Rohinstallationen müssen fertiggestellt sein.'),
    ('5', 'Innenputz ohne Beiputzarbeiten, Estrich und Fassadenarbeiten', Decimal('8.4'), '6, 3 und 3 Prozent der verbleibenden 70 Prozent; alle drei Leistungen müssen fertiggestellt sein.'),
    ('6', 'Fliesenarbeiten im Sanitärbereich und Bezugsfertigkeit, Zug um Zug gegen Besitzübergabe', Decimal('11.2'), '4 und 12 Prozent der verbleibenden 70 Prozent; beide Voraussetzungen müssen vorliegen.'),
    ('7', 'Vollständige Fertigstellung einschließlich protokollierter Restleistungen und Mängelbeseitigung', Decimal('3.5'), '5 Prozent der verbleibenden 70 Prozent.'),
]


def p(doc, text):
    doc.add_paragraph(written_norms(text))


def h(doc, title):
    doc.add_paragraph(title, 'Heading 1')


def table(doc, rows, widths=None):
    add_table(doc, written_norms(rows), widths=widths)


def start(title, *, contract=False):
    d = new_document(title, STAND)
    if contract:
        d.sections[0].left_margin = d.sections[0].right_margin = Cm(2.5)
        d.styles['Normal'].paragraph_format.space_after = Pt(5)
        d.sections[0].header.paragraphs[0].text = 'Notar Dr. Johann Bauer · Hildesheim · Entwurfsfassung'
    p(d, 'Stand 1. November 2027 · Option Einzelverkauf · Entwurf zur Abstimmung')
    return d


def save(d, folder, filename):
    return save_docx(d, OPTION / folder / filename)


def rate_rows(unit):
    price = Decimal(unit['price'])
    return [['Rate', 'Voraussetzung', 'Anteil', 'Betrag EUR']] + [
        [number, description, str(percent).replace('.', ',') + ' %', euro(price * percent / 100)]
        for number, description, percent, _ in RATES
    ]


def rooms(unit):
    if unit['area'] == 68:
        return [('Wohnen und Küche', 26), ('Schlafen', 14), ('Zimmer', 10), ('Bad', 7), ('Flur', 8), ('Abstellraum', 3)]
    if unit['area'] == 82:
        return [('Wohnen und Küche', 29), ('Schlafen', 14), ('Zimmer 1', 11), ('Zimmer 2', 10), ('Bad', 7), ('Flur', 8), ('Abstellraum', 3)]
    if unit['n'] >= 7:
        return [('Wohnen und Küche', 28), ('Schlafen', 14), ('Zimmer', 12), ('Bad', 7), ('Flur', 11), ('Abstellraum', 3)]
    return [('Wohnen und Küche', 29), ('Schlafen', 14), ('Zimmer', 12), ('Bad', 7), ('Flur', 10), ('Abstellraum', 3)]


def contract(unit, *, master=False):
    n = unit['n']; code = f'WE{n:02d}'; price = unit['price']; names = unit['names']
    title = f'Bauträgerkaufvertrag · Wohnung {n:02d}'
    d = start(title, contract=True)
    p(d, 'Urkundenverzeichnis: Nummer wird erst im Beurkundungstermin vergeben. Ort: Hildesheim. '
         'Beurkundungstag: noch nicht vereinbart. Diese Fassung enthält die vorgesehenen Erklärungen; '
         'eine Verlesung, Genehmigung, Unterzeichnung oder Grundbucheinsicht ist damit nicht bestätigt.')
    h(d, '1 Beteiligte und Verfahren')
    p(d, 'Zur gemeinsamen Beurkundung vor Notar Dr. Johann Bauer mit dem Amtssitz Hildesheim sind vorgesehen: '
         'Frau Maren Birk, handelnd ausschließlich als einzelvertretungsberechtigte Geschäftsführerin der '
         'Steinbogen Wohnen GmbH, Wohnhof Am Steinbogen 18, 31134 Hildesheim, als Verkäuferin, sowie ' +
         unit['born'] + ', wohnhaft ' + unit['address'] + ', als Käuferseite. '
         'Die aktuelle Registerbezeichnung und die Vertretungsbefugnis der Verkäuferin sowie die Personalien '
         'werden anhand aktueller Unterlagen im Beurkundungsverfahren festgestellt.')
    p(d, f'Die Käuferseite erwirbt für {unit["purpose"]}. Sie handelt bei diesem Erwerb als Verbraucher. '
         'Eine spätere Vermietung aus dem Privatvermögen ändert diese Einordnung nicht allein wegen '
         'umsatzsteuerrechtlicher Begriffe. Die Beteiligten handeln jeweils im eigenen wirtschaftlichen Interesse. '
         'Etwaige abweichende wirtschaftlich Berechtigte sind dem Notar vor der Beurkundung offenzulegen.')
    p(d, ('Die Eheleute erwerben jeweils einen hälftigen Anteil am nachstehend bezeichneten Wohnungseigentum. '
           'Für den Entwurf ist gesetzlicher Güterstand nach deutschem Recht vorgesehen; ausländische Bezüge, '
           'Eheverträge oder Verfügungsbeschränkungen sind vor Beurkundung festzustellen. '
           'Die Käufer haften für ihre vertraglichen Zahlungspflichten als Gesamtschuldner. '
           if unit['married'] else
           'Der Käufer erwirbt zu Alleineigentum. Für den Entwurf ist angegeben, dass er unverheiratet ist; '
           'der Familienstand und etwaige Verfügungsbeschränkungen sind vor Beurkundung festzustellen. ')+
         'Mitteilungen an mehrere Käufer müssen allen zugehen, soweit nicht nach Vertragsschluss eine gesonderte '
         'Empfangsvollmacht erteilt wird. Eine Empfangsvollmacht für Zahlungsaufforderungen ist kein Verzicht auf Einwendungen.')
    p(d, 'Das Notariat soll der Käuferseite den vollständigen beabsichtigten Vertragstext mit den Anlagen '
         'rechtzeitig, im Regelfall mindestens zwei Wochen vor dem Termin, zur Verfügung stellen. '
         'Der tatsächliche Zugangszeitpunkt wird in der Nebenakte nachgewiesen und erst dann in die Niederschrift übernommen. '
         'Eine Verkürzung wird mit diesem Entwurf weder verlangt noch erklärt. Wesentliche spätere Änderungen werden '
         'mit dem Notariat besprochen, damit ausreichend Gelegenheit zur Prüfung bleibt.')
    h(d, '2 Grundstück, Teilung und Vertragsgegenstand')
    p(d, 'Ausgangspunkt ist das Grundstück ' + GRUNDSTUECK + '. Die Verkäuferin ist im Hauptprojekt als '
         'Grundstückseigentümerin vorgesehen. In Abteilung III besteht die projektbezogene Grundschuld zugunsten der '
         'Leinebogen Projektbank AG über 1.600.000,00 EUR mit dinglichen Zinsen und Nebenleistung gemäß '
         'Grundschuldbestellung UVZ 139/2026 der Notarin Dr. Carla Fink. Die Käuferseite übernimmt diese Belastung '
         'nicht. Maßgeblich für die endgültige Urkunde sind der aktuelle vollständige Grundbuchinhalt und offene Anträge; '
         'der Entwurf behauptet keine zwischenzeitliche Löschung oder aktuelle Einsicht.')
    p(d, f'Die Verkäuferin verkauft der Käuferseite den zu bildenden Miteigentumsanteil von {unit["area"]}/600 '
         f'an diesem Grundstück, verbunden mit dem Sondereigentum an der Wohnung Nummer {n:02d} '
         f'im {unit["floor"].replace("es Obergeschoss", "en Obergeschoss")}, {unit["side"]}, und dem dieser Wohnung zugeordneten Sondernutzungsrecht '
         f'am Pkw-Stellplatz Nummer {n:02d}. Die genaue Abgrenzung ergibt sich aus dem noch behördlich zu '
         'bescheinigenden Aufteilungsplan. Es wird kein Keller, kein Dachterrassenrecht und kein ausschließliches '
         'Gartennutzungsrecht verkauft. Im Wohnungsinneren liegt ein eigener Abstellraum.')
    p(d, f'Die vereinbarte Wohnfläche beträgt {unit["area"]},00 m² nach der beigefügten wohnungsbezogenen '
         'Raumflächenliste. Die aufgeführten Räume werden vollständig und Außenflächen nicht angerechnet. '
         'Eine pauschale Minderflächentoleranz oder ein Verzicht auf Rechte bei Flächenabweichungen ist nicht '
         'vereinbart. Die Festpreisvereinbarung lässt gesetzliche Rechte wegen einer Abweichung von der '
         'vereinbarten Beschaffenheit unberührt. Bewegliche Möbel in Plandarstellungen gehören nicht zum Kaufgegenstand.')
    p(d, 'Die Verkäuferin verpflichtet sich, die Teilung entsprechend dem beigefügten Teilungserklärungsentwurf '
         'und der abgestimmten Aufteilung durchzuführen. Zur Beurkundung müssen ein hinreichend bestimmter '
         'Aufteilungsplan und der vollständige Text der einzubeziehenden Bezugsurkunde vorliegen. Die Anlage '
         '„Vorbereitung Aufteilungsplan“ ist bis zur Ergänzung der behördlich erforderlichen Unterlagen eine '
         'Planungsvorgabe und keine Abgeschlossenheitsbescheinigung. Vor Eintragung des Wohnungseigentums '
         'und der Erwerbervormerkung dürfen keine Kaufpreisraten angefordert werden.')
    h(d, '3 Herstellungspflicht und Beschaffenheit')
    p(d, 'Die Verkäuferin errichtet das Wohngebäude mit acht Wohnungen, drei oberirdischen Geschossen und '
         '820 m² Brutto-Grundfläche ohne Keller auf eigene Verantwortung und Kosten. Sie schuldet die '
         'bezugsfertige und vollständige Herstellung des verkauften Sondereigentums sowie des vertragsgemäßen '
         'Gemeinschaftseigentums einschließlich Erschließung und Außenanlagen. Sie darf geeignete Fachunternehmen '
         'und Planer beauftragen, bleibt aber der Käuferseite für die vertragliche Leistung verantwortlich.')
    p(d, 'Vertragsbestandteile werden die Baubeschreibung BT-BB-01, die Raumflächenliste für diese Wohnung, '
         'die in der Anlagenliste nach Bezeichnung und Stand bestimmten Pläne, der Zahlungsplan, die '
         'Sondernutzungszuordnung und der beurkundete Teilungstext. Tatsächlich ausgehandelte Einzelabreden gehen '
         'vor. Eine unaufgelöste Unstimmigkeit wird vor Beurkundung geklärt und im betroffenen Dokument berichtigt; '
         'die Verkäuferin darf den Leistungsumfang nicht durch freie Wahl des für sie günstigeren Dokuments bestimmen.')
    p(d, 'Die genehmigte Wohnnutzung, die zwingenden öffentlich-rechtlichen Anforderungen und die bei Abnahme '
         'maßgeblichen anerkannten Regeln der Technik bleiben geschuldet. Alle Wohnungen sind barrierefrei '
         'herzustellen; Wohnung 01 wird zusätzlich rollstuhlgerecht mit den konkret vereinbarten Bewegungsflächen '
         'ausgeführt. Für den Schallschutz werden die in BT-BB-01 bezeichneten numerischen Vertragswerte '
         'vereinbart. Ein bloßer Hinweis auf eine Mindestnorm ersetzt diese Beschaffenheitsvereinbarung nicht.')
    p(d, 'Nicht enthalten sind eine Einbauküche, lose Möblierung, individuelle Haushaltsgeräte, eine individuelle '
         'PV-Anlage je Wohnung oder ein garantierter Mieterstromtarif. Die gemeinsame PV-Anlage mit 36 kWp dient '
         'dem vorgesehenen Allgemeinstrom- und Überschusskonzept. Ein steuerlicher Vorteil, eine bestimmte Rendite, '
         'ein Wiederverkaufspreis oder ein bestimmter individueller Verbrauch wird nicht zugesagt. Bewusst falsche '
         'Angaben und ausdrücklich übernommene Zusagen werden durch diese Leistungsabgrenzung nicht entkräftet.')
    h(d, '4 Kaufpreis und enthaltene Kosten')
    p(d, f'Der Kaufpreis beträgt {euro(price)} EUR. Darin enthalten ist das vereinbarte Sondernutzungsrecht '
         'am Stellplatz mit einem kalkulatorischen Teilbetrag von 15.000,00 EUR. Die Verkäuferin erhält den '
         'gesamten Festpreis für Grundstücksanteil, Herstellung und die im Leistungsumfang eingeschlossenen '
         'Anschlüsse. Die Aufgliederung des Stellplatzanteils begründet keine selbstständig abrufbare zusätzliche Rate. '
         'Die Angabe ist keine steuerliche Kaufpreisaufteilung zwischen Grund und Boden und Gebäude.')
    p(d, 'Der Preis umfasst die für das genehmigte Vorhaben erforderlichen Erschließungs- und Hausanschlussleistungen, '
         'die Herstellung der vereinbarten Stellplätze, Fahrradflächen und Spielfläche sowie die Kosten der Teilung '
         'und der hierfür benötigten Unterlagen. Die Verkäuferin trägt öffentliche Beiträge und Anschlusskosten für '
         'die dem Vertrag zugrunde liegende erstmalige Herstellung, auch wenn ein entsprechender Bescheid erst '
         'nach Besitzübergabe ergeht. Nicht eingeschlossen sind nach Besitzübergabe neu beschlossene öffentliche '
         'Verbesserungsmaßnahmen ohne Bezug zur geschuldeten erstmaligen Herstellung.')
    p(d, 'Eine zusätzliche Umsatzsteuer wird der Käuferseite nicht in Rechnung gestellt. Unberührt bleiben '
         'die ausdrücklich vereinbarten Erwerbsnebenkosten und spätere wirksam vereinbarte Sonderwünsche. '
         'Baupreissteigerungen, die normale Beschaffungsorganisation und die Finanzierungskosten der Verkäuferin '
         'berechtigen nicht zu einer einseitigen Kaufpreiserhöhung. Die wirtschaftliche Kalkulation des Projekts '
         'und die Freigabe des alternativen Verkaufszweigs werden außerhalb dieser Urkunde geführt.')
    h(d, '5 Allgemeine Fälligkeitsvoraussetzungen')
    p(d, 'Kaufpreiszahlungen dürfen erst verlangt und entgegengenommen werden, wenn der Vertrag rechtswirksam '
         'ist, die zu seinem Vollzug erforderlichen Genehmigungen vorliegen und der Notar dies schriftlich mitgeteilt '
         'hat. Der Verkäuferin werden keine vertraglichen Rücktrittsvorbehalte eingeräumt. Das verkaufte '
         'Wohnungseigentum muss im Grundbuch begründet und die Vormerkung zur Sicherung des Erwerbsanspruchs '
         'an der vereinbarten Rangstelle eingetragen sein. Nicht übernommene vorrangige oder gleichrangige '
         'Grundpfandrechte müssen nach Maßgabe der MaBV zur Freistellung gesichert sein.')
    p(d, 'Die Freistellungserklärung der Projektbank muss der Käuferseite ausgehändigt werden. Sie muss den '
         'Fall vollständiger und unvollständiger Bauausführung umfassen und den gesetzlich zulässigen Inhalt '
         'haben. Die bloße Benennung einer Ablösesumme oder die Hoffnung auf spätere Bankzustimmung genügt nicht. '
         'Der Grundpfandrechtsgläubiger darf nur nach dem ausgehändigten Freistellungsmechanismus befriedigt '
         'werden. Erwerberzahlungen auf ein ausdrücklich bezeichnetes Ablösekonto wirken im entsprechenden Umfang '
         'gegenüber der Verkäuferin kaufpreiserfüllend.')
    p(d, 'Ferner muss die für die konkrete Ausführung erforderliche Baugenehmigung vorliegen und die '
         'Bauausführung in ihrem zulässigen Umfang gestatten. Der Projektstand sieht die Genehmigung vom '
         '16. Juli 2027 vor; Abweichungen und Nebenbestimmungen sind vor dem tatsächlichen Abruf zu prüfen. '
         'Die notarielle Fälligkeitsmitteilung ersetzt keine Bautenstandsprüfung. Eine Bestätigung der Verkäuferin '
         'über einen Baufortschritt ersetzt ihrerseits weder die Notarmitteilung noch die Bankfreistellung.')
    h(d, '6 Sieben Raten und Fertigstellungssicherheit')
    table(d, rate_rows(unit), [1.4, 9.3, 1.8, 3.5])
    p(d, 'Jede Rate wird frühestens vierzehn Kalendertage nach Zugang einer nachvollziehbaren schriftlichen '
         'Zahlungsanforderung fällig, sofern sämtliche allgemeinen Fälligkeitsvoraussetzungen und sämtliche '
         'in der betreffenden Zeile genannten Bauleistungen erfüllt sind. Bereits vorhandene Voraussetzungen '
         'dürfen in einer Anforderung erläutert werden; die Anforderung begründet keine zusätzliche achte Rate. '
         'Der Baufortschritt ist anhand des zugeordneten Bautenstandsberichts und auf angemessenen Wunsch '
         'durch Besichtigung nachzuweisen. Gesetzliche Zurückbehaltungsrechte bleiben unberührt.')
    p(d, f'Bei der ersten Abschlagszahlung erhält die Käuferseite eine Sicherheit von fünf Prozent des '
         f'Kaufpreises, mithin {euro(Decimal(price) * Decimal("0.05"))} EUR, für die rechtzeitige Herstellung '
         'ohne wesentliche Mängel. Die Verkäuferin verlangt zunächst die Erbringung durch Einbehalt von den '
         'Abschlagszahlungen. Deshalb sind bei Vorliegen sämtlicher Voraussetzungen aus der ersten Rate '
         f'zunächst höchstens {euro(Decimal(price) * Decimal("0.25"))} EUR auszuzahlen. '
         'Die Verkäuferin kann den Einbehalt durch eine den gesetzlichen Anforderungen entsprechende Garantie '
         'oder ein entsprechendes Zahlungsversprechen eines zugelassenen Kreditinstituts oder Kreditversicherers '
         'ablösen. Die Sicherung wird erst freigegeben, wenn ihr Sicherungszweck entfallen ist.')
    p(d, 'Erhöht sich die Vergütung durch Änderungen oder Ergänzungen um mehr als zehn Prozent, ist bei '
         'der nächsten Abschlagszahlung eine zusätzliche Sicherheit von fünf Prozent des zusätzlichen '
         'Vergütungsanspruchs zu leisten. Die Fertigstellungssicherheit ersetzt weder die Pfandfreistellung noch '
         'einen gesetzlichen Einbehalt wegen konkreter Mängel. Eine Bürgschaft nach Paragraf 7 MaBV zur '
         'Abweichung vom vereinbarten Ratenplan wird nicht vereinbart.')
    p(d, 'Die sechste Rate setzt neben den Fliesenarbeiten Bezugsfertigkeit voraus und ist nur Zug um Zug '
         'gegen Besitzübergabe zu zahlen. Bezugsfertigkeit verlangt eine gefahrlose, genehmigungskonforme '
         'Wohnnutzung, funktionierende Ver- und Entsorgung sowie sichere Zugänge und Rettungswege. '
         'Die siebte Rate setzt die vollständige Fertigstellung des vertraglich geschuldeten Objekts sowie die '
         'Erledigung der in den Abnahmeprotokollen bezeichneten Restleistungen und Mängel voraus. '
         'Allein die Abnahme oder die Nutzbarkeit der Wohnung macht diese Schlussrate nicht fällig.')
    h(d, '7 Zahlungsweg, Verzug und Erwerberfinanzierung')
    p(d, 'Zahlungen erfolgen unbar auf die in der notariellen Abwicklung bestätigte, auf die Verkäuferin '
         'lautende Bankverbindung oder das danach zulässige Ablösekonto. Eine Bankverbindung wird erst '
         'nach Abstimmung mit der Projektbank eingesetzt. Änderungen werden über einen zuvor bekannten '
         'Kontaktweg überprüft. Bargeld, Kryptowerte, Gold, Platin und Edelsteine bewirken keine Kaufpreiserfüllung. '
         'Die Beteiligten legen dem Notar die gesetzlich erforderlichen Zahlungsnachweise vor.')
    p(d, 'Kommt die Käuferseite mit einer fälligen und durchsetzbaren Zahlung in Verzug, gelten die '
         'gesetzlichen Verzugszinsen und Schadensersatzregeln. Rechte auf Aufrechnung, Zurückbehaltung '
         'und gesetzliche Leistungsverweigerung werden nicht durch einen pauschalen Ausschluss verkürzt. '
         'Ein Rücktritt der Verkäuferin wegen Zahlungsverzugs setzt die gesetzlichen Voraussetzungen '
         'einschließlich einer erforderlichen Nachfrist voraus. Ein bloßer Finanzierungsengpass ohne '
         'fällige Forderung ist kein vertraglicher Rücktrittsgrund.')
    p(d, 'Die Käuferseite beschafft ihre Finanzierung selbst. Die Verkäuferin gestattet eine vorzeitige '
         'Belastung ausschließlich des verkauften Wohnungseigentums zur Kaufpreisfinanzierung bei einem '
         'in Deutschland zum Geschäftsbetrieb befugten Kreditinstitut, wenn die Grundschuld bis zur '
         'Eigentumsumschreibung nur solche Zahlungen sichert, die mit kaufpreistilgender Wirkung geleistet '
         'werden. Vorherige Verkäuferbelastungen sind nach der abgestimmten Rang- und Freistellungsregelung '
         'zu behandeln. Die Käuferseite übernimmt sämtliche Kosten ihrer Finanzierung und stellt die '
         'Verkäuferin von persönlicher Haftung frei; die Verkäuferin übernimmt kein persönliches Schuldversprechen.')
    p(d, 'Die hierfür nötige Belastungsvollmacht darf nur beim beurkundenden Notar oder dessen amtlicher '
         'Vertretung verwendet werden. Die Finanzierungsgläubigerin muss die Sicherungszweckbeschränkung '
         'und die Auszahlung nach den Fälligkeitsregeln anerkennen. Bei Scheitern des Erwerbs darf sie '
         'eine Löschung nur von der Rückzahlung der mit kaufpreistilgender Wirkung ausgezahlten Beträge '
         'abhängig machen. Die dingliche Haftung wird nach Kaufpreiszahlung beziehungsweise Umschreibung '
         'von der Käuferseite übernommen; die Vollmacht umfasst keine neue Belastung anderer Einheiten.')
    h(d, '8 Termine und Mitwirkung')
    p(d, 'Die Verkäuferin beginnt die Bauarbeiten spätestens am 1. Dezember 2027 und stellt die Wohnung '
         'einschließlich ihrer nutzbaren Erschließung spätestens am 15. September 2028 bezugsfertig her. '
         'Die vollständige Fertigstellung einschließlich Außenanlagen ist spätestens zum 30. September 2028 '
         'geschuldet. Diese Fristen setzen einen spätestens am 30. November 2027 abgeschlossenen Vertrag voraus; '
         'bei einem späteren Abschluss sind die Termine vor Beurkundung ausdrücklich neu zu vereinbaren. '
         'Eine einseitige Terminverschiebung durch die Verkäuferin ist nicht gestattet.')
    p(d, 'Vorhersehbare jahreszeitliche Witterung, gewöhnliche Personalengpässe und die übliche '
         'Materialbeschaffung sind in diesen Fristen berücksichtigt. Nicht von der Verkäuferin zu vertretende '
         'außergewöhnliche Hindernisse werden unverzüglich mit Ursache, betroffenen Arbeiten, Beginn, '
         'voraussichtlicher Dauer und Minderungsmaßnahmen mitgeteilt. Eine Terminfolge setzt eine '
         'tatsächlich ursächliche Behinderung voraus; ihr Umfang richtet sich nach Gesetz und einer konkret '
         'getroffenen Vereinbarung, nicht nach einer pauschalen Fristverlängerung.')
    p(d, 'Bemusterungen erfolgen nach der beigefügten Auswahl- und Terminordnung. Die Verkäuferin '
         'kündigt die erforderliche Auswahl mit angemessener Frist an und benennt die vereinbarte '
         'Basisausstattung. Wenn die Käuferseite trotz Erinnerung und angemessener Nachfrist keine '
         'abweichende Auswahl trifft, wird die bereits in der Baubeschreibung festgelegte Basis ausgeführt. '
         'Damit werden weder ein kostenpflichtiger Sonderwunsch fingiert noch gesetzliche Ansprüche wegen '
         'fehlerhafter oder unklarer Bemusterungsaufforderung ausgeschlossen.')
    h(d, '9 Sonderwünsche und Änderungen')
    p(d, 'Die Verkäuferin ist ohne gesonderte Einigung nicht verpflichtet, Sonderwünsche auszuführen. '
         'Ein Sonderwunsch wird erst verbindlich, wenn Leistung, Vergütungsänderung einschließlich '
         'ersparter Grundleistung, Auswirkungen auf die Termine und gegebenenfalls erforderliche '
         'Genehmigungen vor Ausführung vereinbart sind. Die Beteiligten beachten die gesetzlich '
         'erforderliche Form; eine einfache Nachricht ersetzt eine notwendige notarielle Beurkundung nicht. '
         'Direktaufträge der Käuferseite an am Bau beteiligte Firmen dürfen den Leistungssoll der '
         'Verkäuferin nicht ohne klare Schnittstellenvereinbarung verändern.')
    p(d, 'Eine Änderung der vereinbarten Bauausführung ohne gesonderten Änderungsvertrag ist nur '
         'zulässig, wenn nach Vertragsschluss eine konkret benannte zwingende behördliche Anforderung '
         'entsteht, eine zuvor bei fachgerechter Prüfung nicht erkennbare technische Unvereinbarkeit '
         'beseitigt werden muss oder ein bestimmt vereinbartes Produkt trotz rechtzeitiger zumutbarer '
         'Beschaffung nachweisbar dauerhaft entfällt. Die Änderung muss erforderlich und der Käuferseite '
         'bei Abwägung beider Interessen zumutbar sein; Gebrauchstauglichkeit, vereinbarter Standard, '
         'Wert und wesentliches Erscheinungsbild dürfen nicht vermindert werden. Reine Kostenersparnis '
         'und gewöhnliche Lieferverspätung genügen nicht.')
    p(d, 'Vor Durchführung einer solchen Änderung teilt die Verkäuferin den belegten Grund, das '
         'Ersatzprodukt oder Detail und die Auswirkungen auf Gebrauch, Gestaltung, Wartung und Kosten mit. '
         'Der Kaufpreis erhöht sich dadurch nicht. Die Wohnungsabgrenzung, die zugesagte Wohnfläche, '
         'die Barrierefreiheit, der Schallschutz und der zugeordnete Stellplatz dürfen nicht einseitig '
         'verschlechtert werden. Für nicht von diesem eng begrenzten Vorbehalt erfasste Abweichungen '
         'ist eine ausdrückliche Einigung in der erforderlichen Form notwendig.')
    p(d, 'Eine Änderung der Teilungserklärung kann nur verlangt werden, soweit sie konkret der Erfüllung '
         'einer nach Vertragsschluss entstandenen behördlichen Auflage oder der Beseitigung eines '
         'nachträglich erkannten zeichnerischen Zuordnungsfehlers dient und das Sondereigentum, '
         'Sondernutzungsrechte, Stimmgewicht, Kostenanteile und der vertragsgemäße Gebrauch der '
         'Käuferseite nicht beeinträchtigt werden. Die Verkäuferin legt den vollständigen Änderungstext '
         'und die Begründung zur Prüfung vor. Eine allgemeine Vollmacht zur wirtschaftlichen Umgestaltung '
         'oder eine Zustimmung zu beliebigen Nachträgen wird nicht erteilt. Zusätzliche Wohnungen, '
         'Gewerbenutzungen und die Umwandlung des Wartungsdachs in eine Dachterrasse sind nicht umfasst.')
    h(d, '10 Baustellenzutritt und Dokumentation')
    p(d, 'Die Verkäuferin ermöglicht nach angemessener Voranmeldung begleitete Besichtigungen des '
         'Kaufobjekts und die Vorbereitung der Abnahme mit einer sachkundigen Vertrauensperson. '
         'Sie darf den Zutritt aus konkreten Gründen der Arbeitssicherheit zeitlich koordinieren und '
         'einen zeitnahen Ersatztermin anbieten. Ein pauschales Zutrittsverbot oder ein Verbot sachverständiger '
         'Begleitung wird nicht vereinbart. Weisungen an Unternehmer erfolgen über die Bauleitung, '
         'damit die technische Verantwortung und der Bauablauf eindeutig bleiben.')
    p(d, 'Die Verkäuferin übergibt spätestens mit der vollständigen Fertigstellung die für Betrieb '
         'und Instandhaltung erforderlichen Bestandsunterlagen, Bedienungs- und Wartungsanleitungen, '
         'Prüf- und Inbetriebnahmeberichte, wohnungsbezogene Zählerzuordnung und den Gebäudeenergieausweis '
         'nach den gesetzlichen Vorgaben. Gemeinschaftliche Unterlagen erhält zusätzlich die Verwaltung '
         'für die GdWE. Ein Bautagebuchauszug oder eine Fachunternehmerbescheinigung ersetzt nicht '
         'fehlende vertraglich geschuldete Pläne oder technische Nachweise.')
    h(d, '11 Abnahme von Sonder- und Gemeinschaftseigentum')
    p(d, 'Die Abnahme erfolgt nach Herstellung der Abnahmereife in einer gemeinsamen Besichtigung '
         'mit Niederschrift. Die Verkäuferin lädt mit mindestens vierzehn Kalendertagen Vorlauf ein '
         'und bezeichnet den Gegenstand sowie die verfügbaren Unterlagen. Die Käuferseite darf eine '
         'sachkundige Person ihrer Wahl hinzuziehen. Die Kosten einer von der Verkäuferin eingesetzten '
         'Abnahmeorganisation trägt diese; eine von der Käuferseite selbst beauftragte zusätzliche '
         'Begutachtung richtet sich nach deren gesondertem Auftrag und etwaigen gesetzlichen Ersatzansprüchen.')
    p(d, 'Die Käuferseite erklärt die Abnahme des Sondereigentums und des auf ihren Erwerbsvertrag '
         'bezogenen Gemeinschaftseigentums selbst. Eine gemeinsame Begehung mehrerer Erwerber '
         'ersetzt die einzelne Erklärung nicht. Der Verwalter, ein von der Verkäuferin beauftragter '
         'Sachverständiger, eine Erwerbermehrheit und drei ausgewählte Erwerber erhalten durch '
         'diesen Vertrag keine Abnahmevollmacht. Eine später frei erteilte, inhaltlich bestimmte '
         'Vollmacht an eine Vertrauensperson bleibt möglich. Ein früheres Abnahmeprotokoll anderer '
         'Erwerber bindet diese Käuferseite nicht.')
    p(d, 'Das Protokoll hält die untersuchten Bereiche, festgestellte Mängel, noch ausstehende '
         'Leistungen, die konkrete Abnahmeerklärung oder ihre begründete Verweigerung und ausdrücklich '
         'erklärte Vorbehalte fest. Eine Abnahme darf wegen unwesentlicher Mängel nicht verweigert '
         'werden. Rechte wegen bekannter Mängel werden bei Abnahme nach Maßgabe des Gesetzes '
         'vorbehalten. Besitzübergabe, Zahlung, Schweigen oder Einzug gelten nicht aufgrund dieses '
         'Vertrags als Abnahme. Die gesetzliche Regelung des Paragrafen 640 Absatz 2 BGB bleibt '
         'mit sämtlichen Voraussetzungen, insbesondere dem erforderlichen Verbraucherhinweis, unberührt.')
    h(d, '12 Besitz, Nutzen, Lasten und Gefahr')
    p(d, 'Die Verkäuferin übergibt die bezugsfertige Wohnung geräumt, besenrein und frei von '
         'Miet- oder sonstigen Nutzungsrechten, sobald die bis dahin fälligen und durchsetzbaren '
         'Kaufpreisraten erfüllt sind; die sechste Rate ist Zug um Zug gegen Übergabe zu leisten. '
         'Ein berechtigter Mängeleinbehalt verhindert die Übergabe nicht allein deshalb, weil '
         'dadurch noch ein Betrag offensteht. Schlüssel, Zählerstände und Dokumente werden im '
         'Übergabeprotokoll erfasst. Die Übergabe ist von der Abnahmeerklärung und der Eigentumseintragung zu unterscheiden.')
    p(d, 'Mit Besitzübergabe gehen Nutzungen, laufende öffentliche und private Lasten sowie '
         'die tatsächliche Verkehrssicherung für den übergebenen Bereich auf die Käuferseite '
         'über. Die gesetzlichen Regeln zur Gefahrtragung und zu von der Verkäuferin zu '
         'vertretenden Schäden bleiben unberührt. Gemeinschaftliche Versicherungs-, Betriebs- '
         'und Verwaltungskosten werden ab diesem Zeitpunkt nach der wirksamen Gemeinschaftsordnung '
         'und den gesetzlichen Vorschriften abgerechnet. Die Verkäuferin trägt die auf noch '
         'nicht übergebene Einheiten entfallenden Kosten selbst.')
    p(d, 'Eine Vermietung durch die Verkäuferin vor der vertraglich geschuldeten lastenfreien '
         'Übergabe ist ausgeschlossen. Beabsichtigt die Käuferseite eine Vermietung, schließt '
         'sie eigene Mietverträge und kalkuliert die daraus folgenden Pflichten selbst. '
         'Mietgarantie, Vermietungsauftrag und garantierte Erträge sind nicht vereinbart. '
         'Soll das Hauptprojekt bereits vermietet worden sein, ist diese unvermietete '
         'Verkaufsoption vor Vertragsschluss vollständig an den tatsächlichen Bestand anzupassen.')
    h(d, '13 Mängelrechte und Nacherfüllung')
    p(d, 'Für die Bauleistung gelten die gesetzlichen Mängelrechte. Die Verjährung bei '
         'Bauwerksmängeln beträgt im gesetzlichen Grundfall fünf Jahre ab wirksamer '
         'Abnahme des jeweils betroffenen Leistungsgegenstands. Unterschiedliche '
         'Abnahmezeitpunkte von Sonder- und Gemeinschaftseigentum sind getrennt festzuhalten. '
         'Arglist, übernommene Garantien, Hemmung, Neubeginn und sonstige gesetzliche '
         'Sonderregeln bleiben unberührt. Ein vertraglicher Ausschluss neuer Bauwerksmängel '
         'oder eine Bindung an das Verjährungsende anderer Erwerber wird nicht vereinbart.')
    p(d, 'Die Käuferseite soll Mängel mit Ort, Erscheinung, Feststellungszeitpunkt und '
         'gegebenenfalls vorhandenen Bildern anzeigen und der Verkäuferin Gelegenheit '
         'zur Prüfung und gesetzlich geschuldeten Nacherfüllung geben. Ein bestimmtes '
         'Portal oder eine besondere Form ist keine Ausschlussvoraussetzung. Die '
         'Verkäuferin stimmt notwendige Besichtigung und Arbeiten rechtzeitig ab. '
         'Bei Gefahr im Verzug und in den übrigen gesetzlich geregelten Fällen bleiben '
         'Sofortmaßnahmen, Selbstvornahme und Kostenerstattung möglich.')
    p(d, 'Die Verkäuferin darf die ordnungsgemäße Art der Nacherfüllung im gesetzlichen '
         'Rahmen wählen. Sie schuldet keinen vollständigen Austausch mangelfreier '
         'Bauteile, wenn eine fachgerechte gleichwertige Mangelbeseitigung den '
         'vertraglichen Zustand herstellt und die gesetzlichen Voraussetzungen erfüllt. '
         'Verschleiß, fehlende vereinbarte Wartung und nachträgliche Eingriffe sind '
         'ursachenbezogen zu prüfen; sie führen nicht automatisch zum Verlust '
         'sämtlicher Gewährleistungsrechte. Die Abtretung von Ansprüchen gegen '
         'Nachunternehmer ersetzt die Haftung der Verkäuferin nicht.')
    p(d, 'Soweit gesetzlich zulässig, kann die GdWE gleichgerichtete Ansprüche wegen '
         'Mängeln am Gemeinschaftseigentum durch Beschluss zur gemeinschaftlichen '
         'Verfolgung übernehmen. Dieser Vertrag schließt solche Beschlüsse nicht aus '
         'und begründet keine zwingende Abtretung aller Ansprüche. Eine doppelte '
         'Befriedigung desselben Mangelbeseitigungsinteresses kann nicht verlangt werden. '
         'Die Anzeige und Bearbeitung gemeinschaftlicher Mängel ist mit der Verwaltung '
         'zu koordinieren, ohne eine eigenständige Ausschlussfrist zu schaffen.')
    h(d, '14 Auflassung, Vormerkung und Vollzug')
    p(d, f'Die Beteiligten sind darüber einig, dass das in Ziffer 2 bezeichnete '
         f'Wohnungseigentum auf {names} '+('zu je einem hälftigen Anteil' if unit['married'] else 'zu Alleineigentum')+
         ' übergehen soll. Die Auflassung wird in der endgültigen Urkunde unbedingt erklärt. '
         'Die Verkäuferin bewilligt und die Käuferseite beantragt die Eintragung '
         'einer Vormerkung zur Sicherung des Anspruchs auf Eigentumsübertragung '
         'an der vereinbarten Rangstelle. Ihre endgültige Bezeichnung wird an '
         'die wirksam angelegten Wohnungsgrundbücher angepasst; eine nicht existente '
         'Blattnummer wird im Entwurf nicht als Registertatsache ausgegeben.')
    p(d, 'Der Notar soll die Umschreibung beantragen, sobald die vertraglich '
         'geschuldete Gegenleistung bezahlt oder ein verbleibender Streitbetrag '
         'durch übereinstimmende notarielle Abwicklungsanweisung oder gerichtliche '
         'Entscheidung geklärt ist, die erforderliche steuerliche '
         'Unbedenklichkeitsbescheinigung vorliegt und die gesetzlichen '
         'Zahlungsnachweispflichten erfüllt sind. Die gesetzlichen Ansprüche '
         'auf Eigentumsübertragung und deren gerichtliche Durchsetzung werden '
         'dadurch nicht ausgeschlossen. Bis dahin gibt der Notar keine '
         'Ausfertigung mit ungesicherter Vollzugsmöglichkeit heraus.')
    p(d, 'Die Käuferseite bewilligt die Löschung der Vormerkung bei ihrer '
         'Eigentumseintragung, sofern keine beeinträchtigenden Zwischenrechte '
         'ohne ihre Zustimmung bestehen bleiben. Für den Fall des Scheiterns '
         'des Vertrags wird keine automatische Löschungsvollmacht allein '
         'auf einseitige Behauptung der Verkäuferin erteilt. Rückabwicklung, '
         'Rückzahlung empfangener Beträge und Löschung sind dann gesondert '
         'rechtlich und notariell zu sichern.')
    h(d, '15 Vollzugsvollmacht und Gemeinschaftsordnung')
    p(d, 'Die Beteiligten bevollmächtigen den beurkundenden Notar und seine '
         'amtliche Vertretung, Anträge zum vertragsgemäßen Grundbuchvollzug '
         'getrennt oder beschränkt zu stellen und offensichtliche Schreibfehler '
         'oder grundbuchliche Bezeichnungsfehler zu berichtigen. Änderungen '
         'des wirtschaftlichen Vertragsinhalts, des Kaufpreises, der Fläche, '
         'des Leistungsumfangs und der zugeordneten Rechte sind nicht umfasst. '
         'Die Vollmacht berechtigt nicht zur Abnahme, zu Mängelverzichten, '
         'Vergleichen oder der Unterwerfung unter die Zwangsvollstreckung.')
    p(d, 'Die Käuferseite tritt mit dem gesetzlich maßgeblichen Zeitpunkt '
         'in die GdWE ein. Inhalt und Reichweite der Gemeinschaftsordnung '
         'ergeben sich aus dem beurkundeten Teilungstext. Die Verwaltung '
         'kann organisatorische Aufgaben übernehmen, aber keine '
         'vertragliche Abnahme für die Käuferseite erklären. Eine '
         'Erstverwalterbestellung ist in dieser Urkunde nicht enthalten; '
         'sie bedarf einer gesonderten wirksamen Entscheidung. Die '
         'Verkäuferin erhält weder ein dauerhaftes Mehrstimmrecht '
         'noch eine Befreiung von Kosten für ihre unverkauften Einheiten.')
    h(d, '16 Kosten, Belehrungen und Schlussbestimmungen')
    p(d, 'Die Käuferseite trägt die Kosten dieses Erwerbsvertrags und seines '
         'üblichen Vollzugs, der eigenen Vormerkung und Eigentumsumschreibung '
         'sowie die Grunderwerbsteuer. Die Verkäuferin trägt die Kosten '
         'der Teilung, der Abgeschlossenheitsunterlagen und der Freistellung '
         'von ihren nicht übernommenen Belastungen. Jede Seite trägt '
         'die Kosten ihrer eigenen Beratung. Die gesetzliche Außenhaftung '
         'für Steuern und Gebühren bleibt unberührt. Eine selbstständige '
         'Maklerprovisionszusage wird in diesem Vertrag nicht begründet.')
    p(d, 'Der Notar wird über Formbedürftigkeit, Vormerkung, Eigentumserwerb '
         'durch Eintragung, MaBV-Sicherungen, Gemeinschaftseigentum, '
         'Zahlungsnachweise und wirtschaftliche Grenzen seiner Prüfung '
         'belehren. Nebenabreden über Kaufpreis, Bauleistung oder '
         'Sonderwünsche sind vollständig offenzulegen und in der '
         'erforderlichen Form aufzunehmen. Eine tatsächliche Individualabrede '
         'wird nicht durch eine pauschale Vollständigkeitsklausel verdrängt. '
         'Die Unwirksamkeit einer einzelnen Regel richtet sich nach '
         'den gesetzlichen Vorschriften; eine automatische Ersetzung '
         'durch die verkäuferfreundlichste ähnliche Klausel wird nicht vereinbart.')
    h(d, '17 Anlagen und vorgesehener Urkundsabschluss')
    table(d, [['Anlage', 'Gegenstand'], ['1', 'Teilungserklärung und Gemeinschaftsordnung TE-OPTION-01'],
        ['2', 'Baubeschreibung BT-BB-01 mit Planverzeichnis'], ['3', f'Raum- und Zuordnungsliste {code}'],
        ['4', f'Wohnungsbezogener Ratenplan {code}'], ['5', 'Bemusterungs- und Sonderwunschordnung BT-BM-01']], [1.5,12.5])
    p(d, 'Der Schlussvermerk über vollständiges Vorlesen, Vorlage der '
         'Pläne, Genehmigung und eigenhändige Unterzeichnung wird erst '
         'nach tatsächlichem Ablauf der Beurkundungsverhandlung ausgefüllt. '
         'Diese Entwurfsfassung ist weder eine Urschrift noch eine '
         'beglaubigte Abschrift und enthält keine Unterschrift oder Siegelnachbildung.')
    p(d, 'Unterschriftszeile Verkäuferin: __________________________________')
    p(d, 'Unterschriftszeile Käuferseite: __________________________________')
    p(d, 'Unterschriftszeile Notar: ________________________________________')
    if master:
        # Inhaltlich vollständige Vorlage, nicht ein bloßes Klauselskelett.
        d.core_properties.title = 'Wordvorlage Bauträgerkaufvertrag mit vollständigen Vertragsklauseln'
        d.paragraphs[0].text = 'Wordvorlage · Bauträgerkaufvertrag WE01'
        replacements = {names: '[Namen der Käuferseite]', unit['address']: '[Anschrift der Käuferseite]',
                        unit['born']: '[Personalien und Geburtsdaten der Käuferseite]'}
        for paragraph in d.paragraphs:
            for old, new in replacements.items():
                if old in paragraph.text:
                    paragraph.text = paragraph.text.replace(old, new)
        save_docx(d, '12_Wordvorlagen/Bautraeger/V01_Bautraegerkaufvertrag_ausformulierte_Wordvorlage.docx')
    else:
        save(d, '01_Vertragsentwuerfe', f'BT{n:02d}_Bautraegerkaufvertrag_{code}.docx')


def teilung():
    d = start('Teilungserklärung und Gemeinschaftsordnung · TE-OPTION-01', contract=True)
    p(d, 'Dieser Text ist die vorgesehene Bezugsurkunde für einen noch nicht beschlossenen Einzelverkauf. '
         'Es wird keine bereits erklärte Teilung, eingetragene Gemeinschaftsordnung oder behördliche '
         'Abgeschlossenheitsbescheinigung bestätigt. Urkundenverzeichnis und Beurkundungstag werden '
         'erst bei Durchführung des gesonderten Verfahrens vergeben.')
    h(d, '1 Teilende Eigentümerin und Grundbuch')
    p(d, 'Die Steinbogen Wohnen GmbH, vertreten durch Maren Birk als Geschäftsführerin, erklärt '
         'gegenüber dem Grundbuchamt die nachstehende Aufteilung des Grundstücks '+GRUNDSTUECK+
         '. Die aktuelle Eintragung der Eigentümerin und die vorhandenen Belastungen sind vor '
         'Beurkundung durch den Notar festzustellen. Die Projektgrundschuld der Leinebogen Projektbank AG '
         'über 1.600.000,00 EUR wird durch die Aufteilung nicht gelöscht. Für Einzelverkäufe sind '
         'gesonderte Freistellungserklärungen erforderlich.')
    h(d, '2 Bildung der acht Wohnungseigentumsrechte')
    p(d, 'Die Eigentümerin teilt das Grundstück gemäß Paragraf 8 WEG in die nachstehenden '
         'Miteigentumsanteile auf. Mit jedem Anteil wird Sondereigentum an den im endgültigen '
         'Aufteilungsplan mit derselben Nummer gekennzeichneten Wohnungsräumen verbunden. '
         'Die Bruchteile sind bewusst auf den Nenner 600 bezogen und ergeben zusammen 600/600. '
         'Eine ungenaue Rundung auf Tausendstel erfolgt nicht.')
    table(d, [['WE', 'Geschoss', 'Wohnfläche', 'Miteigentum', 'Stellplatz-SNR']] + [
        [f'{u["n"]:02d}', u['floor'], f'{u["area"]},00 m²', f'{u["area"]}/600', f'{u["n"]:02d}'] for u in UNITS],
        [1.1,4.3,2.8,2.8,3.0])
    p(d, 'Das jeweilige Wohnungsgrundbuch erhält seine amtliche Blattnummer erst durch '
         'Anlegung. Der Entwurf reserviert keine erfundenen Blattnummern. Maßgeblich '
         'sind die behördlich bescheinigten Pläne und die Eintragungsbewilligung. '
         'Die geometrische Abgrenzung darf nicht allein durch Wohnflächenzahlen oder '
         'eine informelle Planskizze ersetzt werden.')
    h(d, '3 Umfang des Sonder- und Gemeinschaftseigentums')
    p(d, 'Zum Sondereigentum gehören die rechtlich sondereigentumsfähigen Bestandteile '
         'innerhalb der bezeichneten Wohnungsräume, insbesondere nichttragende '
         'Innenoberflächen und Ausstattungen, soweit Bestand, Sicherheit, '
         'gemeinschaftlicher Gebrauch und Rechte anderer Eigentümer nicht betroffen sind. '
         'Zwingendes Gemeinschaftseigentum wird durch diese Beschreibung nicht '
         'in Sondereigentum umgewandelt.')
    p(d, 'Gemeinschaftlich bleiben insbesondere Grundstück, Fundament und Bodenplatte, '
         'tragende Wände und Decken, Dach und Abdichtung, Fassaden, außenwirksame '
         'Fensterbestandteile, Hauseingang, Treppenraum, Aufzug, gemeinschaftliche '
         'Leitungen, Wärmeerzeugung und Regelung, gemeinsame Lüftungseinrichtungen, '
         'PV-Anlage, Hauptverteilungen und Erschließungsleitungen. Gleiches gilt '
         'für Bauteile, deren gemeinschaftlicher Charakter zwingend aus Paragraf 5 WEG folgt, '
         'auch wenn sie sich räumlich innerhalb einer Wohnung befinden.')
    p(d, 'Das Haus hat keinen Keller. Abstellräume innerhalb der Wohnungen gehören '
         'zur jeweiligen Wohnung; Technik-, Fahrrad- und Außenflächen bleiben '
         'gemeinschaftlich, soweit nachfolgend ein Sondernutzungsrecht zugeordnet '
         'wird. Die zurückgesetzte südliche Dachfläche ist Wartungsdach und '
         'keine gemeinschaftliche oder private Dachterrasse. Nutzung zu '
         'Aufenthaltszwecken ist nicht Vertragsinhalt.')
    h(d, '4 Sondernutzungsrechte an den Stellplätzen')
    p(d, 'Jede Wohnung erhält das ausschließliche Sondernutzungsrecht an genau '
         'dem nummerngleichen Pkw-Stellplatz. Lage, Umgrenzung, Breite und Tiefe '
         'werden im endgültigen Aufteilungsplan dargestellt. Stellplatz 01 '
         'wird mit der für die rollstuhlgerechte Wohnung vorgesehenen barrierefreien '
         'Wegekette verbunden. Das Sondernutzungsrecht erlaubt das Abstellen '
         'eines zulässigen Pkw, aber keine Einhausung, Lagerung gefährlicher '
         'Stoffe oder Behinderung gemeinschaftlicher Wege.')
    p(d, 'Die Oberflächen und ihre konstruktiven Bestandteile bleiben '
         'gemeinschaftlich. Der jeweilige Berechtigte trägt die laufende '
         'Reinigung seines Stellplatzes und die Kosten seiner individuellen '
         'Stromnutzung. Bauliche Erhaltung und Veränderungen erfolgen durch '
         'die GdWE nach Gesetz und wirksamen Beschlüssen. Ein Sondernutzungsrecht '
         'allein erlaubt keinen Eingriff in Abdichtung, Leitungsführung oder '
         'gemeinsame elektrische Leistung. Die erstmalig vereinbarte '
         'Herstellung schuldet weiterhin die Verkäuferin.')
    h(d, '5 Wohnnutzung und gemeinschaftliche Flächen')
    p(d, 'Die Einheiten dienen dem Wohnen. Eine nicht störende Tätigkeit '
         'innerhalb der Wohnung bleibt im gesetzlichen Rahmen zulässig; '
         'eine gewerbliche Umnutzung mit verändertem Publikumsverkehr oder '
         'bauordnungsrechtlichen Anforderungen bedarf der erforderlichen '
         'Genehmigungen und Zustimmungen. Die Vermietung zu Wohnzwecken '
         'ist zulässig. Diese Gemeinschaftsordnung enthält keinen Anspruch '
         'auf baurechtlich unzulässige Nutzungen.')
    p(d, 'Die 16 Bewohner-Fahrradplätze werden so organisiert, dass jede '
         'Wohnung zwei Plätze nutzen kann; die vier Besucherplätze bleiben '
         'allgemein zugänglich. Die 30 m² große Spielfläche und die '
         'Grünflächen dienen der gemeinschaftlichen Nutzung. Die GdWE '
         'kann eine angemessene Benutzungsordnung beschließen. '
         'Rettungswege, Aufstellflächen, Technikzugänge und der barrierefreie '
         'Zugang dürfen nicht zugestellt werden.')
    h(d, '6 Technische Anlagen und Betrieb')
    p(d, 'Die Wärmepumpenkaskade, der Aufzug, die PV-Anlage und das '
         'Lastmanagement werden gemeinschaftlich betrieben. Die '
         'GdWE übernimmt nach ordnungsgemäßer Übergabe die '
         'gesetzlichen Betreiberpflichten und beauftragt die '
         'erforderlichen Wartungen und Prüfungen. Die Übergabe '
         'enthält Anlagenidentifikation, Bedienunterlagen und '
         'Prüfberichte. Vor einer ordnungsgemäßen betriebsfähigen '
         'Herstellung wird der Verkäuferin durch diese Regelung '
         'keine Betreiberpflicht abgenommen.')
    p(d, 'Die PV-Anlage bleibt Gemeinschaftseigentum; Strom wird '
         'für die gemeinschaftlichen Verbraucher eingesetzt und '
         'im Übrigen nach dem gewählten Netz- und Messkonzept '
         'eingespeist. Ein Stromliefervertrag mit einzelnen '
         'Wohnungseigentümern wird hierdurch nicht geschlossen. '
         'Änderungen des Betreibermodells und die damit '
         'verbundenen steuerlichen und energierechtlichen '
         'Pflichten bedürfen einer gesonderten Prüfung und Entscheidung.')
    h(d, '7 Kosten und Lasten')
    p(d, 'Die laufenden gemeinschaftlichen Kosten und Lasten '
         'werden grundsätzlich im Verhältnis der Miteigentumsanteile '
         'verteilt, soweit zwingendes Recht oder wirksame '
         'Vereinbarungen und Beschlüsse eine andere Zuordnung '
         'verlangen. Heiz- und Warmwasserkosten werden nach '
         'den dafür geltenden gesetzlichen Vorgaben abgerechnet. '
         'Individuell erfasster Wohnungsstrom und eigene '
         'Ladevorgänge sind keine unbesehen nach Fläche '
         'zu verteilenden Allgemeinkosten.')
    p(d, 'Die Eigentümerversammlung beschließt Wirtschaftsplan, '
         'Vorschüsse, Rücklage und Jahresabrechnung im '
         'gesetzlichen Verfahren. Diese Urkunde setzt '
         'keine erfundene dauerhafte Hausgeldhöhe fest. '
         'Die Verkäuferin bleibt für ihre Einheiten '
         'kostenpflichtig; ein Leerstand oder fehlender '
         'Verkauf befreit sie nicht von den auf sie '
         'entfallenden gemeinschaftlichen Lasten.')
    h(d, '8 Verwaltung, Stimmrecht und Vertretung')
    p(d, 'Die Eigentümer entscheiden über die Bestellung '
         'eines geeigneten Verwalters in einer gesonderten '
         'Beschlussfassung. Die gesetzlichen Grenzen '
         'für Erstbestellung, Abberufung und Vertragsdauer '
         'bleiben maßgeblich. Dieser Entwurf enthält '
         'keine Bestellung einer mit der Verkäuferin '
         'verbundenen Verwaltung und keinen bereits '
         'geschlossenen Verwaltervertrag.')
    p(d, 'Für das Stimmrecht gilt die gesetzliche Regelung, '
         'soweit nicht im endgültigen notariellen Verfahren '
         'eine ausdrücklich bezeichnete wirksame Abweichung '
         'vereinbart wird. Die Verkäuferin erhält kein '
         'Stimmrecht für bereits auf Erwerber übergegangene '
         'Positionen entgegen Paragraf 8 Absatz 3 WEG. '
         'Der Verwalter vertritt die GdWE im gesetzlichen '
         'Rahmen, aber nicht die Erwerber bei der '
         'kaufvertraglichen Abnahme.')
    h(d, '9 Mängel, Abnahme und Zugang')
    p(d, 'Die kaufvertragliche Abnahme und die aus den '
         'Erwerbsverträgen folgenden Mängelrechte werden '
         'nicht durch diese Gemeinschaftsordnung '
         'verändert. Weder ein Mehrheitsbeschluss noch '
         'eine technische Begehung erklärt ohne '
         'individuelle Rechtsgrundlage die Abnahme '
         'für einen Erwerber. Die gesetzlich mögliche '
         'gemeinschaftliche Verfolgung gleichgerichteter '
         'Mängelansprüche bleibt zulässig.')
    p(d, 'Notwendiger Zugang zu gemeinschaftlichen '
         'Bauteilen in einer Wohnung ist nach '
         'rechtzeitiger Ankündigung und Abstimmung '
         'im gesetzlichen Umfang zu gewähren. '
         'Bei dringender Gefahr gelten die '
         'gesetzlichen Notmaßnahmen. Die dadurch '
         'verursachten Beeinträchtigungen und '
         'Wiederherstellungspflichten werden '
         'ursachenbezogen nach Gesetz behandelt; '
         'ein pauschaler entschädigungsloser '
         'Eingriff in beliebige Wohnungsbereiche '
         'wird nicht vereinbart.')
    h(d, '10 Änderung und Grundbuchvollzug')
    p(d, 'Die Eigentümerin bewilligt und beantragt die '
         'Eintragung der bezeichneten Aufteilung und '
         'Gemeinschaftsordnung unter Bezugnahme auf '
         'den endgültigen behördlichen Aufteilungsplan. '
         'Die Urkunde darf erst mit den gesetzlich '
         'notwendigen Anlagen zum Grundbuchvollzug '
         'eingereicht werden. Die notariellen '
         'Vollzugsbefugnisse sind auf formale '
         'Korrekturen beschränkt; wirtschaftliche '
         'Änderungen bedürfen einer gesonderten '
         'wirksamen Erklärung der Berechtigten.')
    p(d, 'Die Verkäuferin behält sich keine freie '
         'Änderung von Wohnungszahl, Kostenquote, '
         'Nutzungsart, Fläche oder Sondernutzungsrechten '
         'vor. Etwaige Mitwirkungspflichten aus '
         'einem Erwerbsvertrag setzen dessen '
         'konkrete, zumutbare und wirksame '
         'Regelung voraus. Eine unwirksame '
         'Klausel wird nicht durch eine '
         'pauschale Treuepflicht ersetzt.')
    save(d, '02_Gemeinsame_Anlagen', 'AN01_Teilungserklaerung_Gemeinschaftsordnung_Entwurf.docx')


def baubeschreibung():
    d = start('Baubeschreibung BT-BB-01 · Einzelverkaufsoption')
    p(d, 'Die nachstehende Beschreibung konkretisiert das vorgesehene '
         'Leistungssoll der Verkaufsoption. Sie ergänzt den vorhandenen '
         'Planungsstand des Mietprojekts. Ihre Einbeziehung in einen '
         'Bauträgervertrag setzt eine ausdrückliche Entscheidung '
         'für die Option und die vollständige notarielle Abstimmung voraus.')
    sections = [
        ('1 Grundstück und Erschließung', [
            'Das Gebäude wird auf dem 1.300 m² großen Grundstück Wohnhof Am Steinbogen 18 in Hildesheim errichtet. Die äußeren Abmessungen betragen in Erdgeschoss und erstem Obergeschoss 20,00 m mal 15,00 m, im zweiten Obergeschoss 20,00 m mal 11,00 m. Daraus folgen 820 m² Brutto-Grundfläche. Der obere Rücksprung liegt an der Südseite. Ein Keller wird nicht gebaut.',
            'Die Verkäuferin stellt die für den Betrieb erforderlichen Anschlüsse für Wasser, Abwasser, Strom und die vorgesehenen Telekommunikations-Leerrohre her. Die Einführung ins Gebäude wird gegen die zu erwartenden Einwirkungen abgedichtet. Die Außenentwässerung und Rückhaltung entsprechen dem genehmigten, koordinierten Entwässerungsplan. Die öffentliche Infrastruktur wird nicht als Eigentum der Erwerber verkauft.',
        ]),
        ('2 Tragwerk, Bodenplatte und Rohbau', [
            'Der Rohbau erhält ein auf den geotechnischen Untersuchungsstand abgestimmtes Gründungssystem mit Bodenplatte und verdichteter Tragschicht. Tragende Bauteile werden nach der abschließend geprüften beziehungsweise verfahrensgerecht erstellten Statik hergestellt. Die koordinierte Öffnungsplanung wird vor Ausführung mit Tragwerk und Haustechnik abgeglichen. Nachträgliche Kernbohrungen durch bewehrte oder brandschutzrelevante Bauteile dürfen keine ungenehmigten Ersatzlösungen für fehlende Koordination bilden.',
            'Außen- und Wohnungstrennwände, Decken und Treppen werden einschließlich der für Schall-, Brand- und Wärmeschutz erforderlichen Anschlüsse ausgeführt. Die Verkäuferin schuldet ein zusammenpassendes Bausystem; eine isolierte Baustoffkennzahl ersetzt nicht die Funktionsfähigkeit der Anschlüsse. Estrich und Bodenbeläge erhalten die abgestimmten Rand- und Bewegungsfugen.',
        ]),
        ('3 Dach, Fassade und Abdichtung', [
            'Die Dachflächen werden als abgedichtete Flachdachkonstruktion mit geregelter Haupt- und Notentwässerung hergestellt. Das obere Dach liegt bei +9,60 m, die Attika endet bei +9,90 m über dem Bezugsgelände. Die südliche Rücksprungfläche liegt bei +6,20 m und bleibt ausschließlich Wartungsdach. PV-Unterkonstruktion, Dachabläufe, Anschlüsse und Wartungswege sind konstruktiv zu koordinieren.',
            'Die Fassade wird als fertig beschichtete, gedämmte Gebäudehülle mit zum System passenden Sockel-, Fensterbank- und Dachrandanschlüssen ausgeführt. Die äußere Farbgebung folgt der neutralen hellen Grundpalette des abgestimmten Fassadenentwurfs. Eine nachträgliche beliebige Farbänderung durch die Verkäuferin ist nicht vereinbart. Fensteranschlüsse werden innen luftdicht und außen schlagregensicher sowie entwässerungsgerecht hergestellt.',
            'Der Hauseingang wird schwellenlos mit der freigegebenen Entwässerungsrinne und einem vom Gebäude weg gerichteten Oberflächengefälle ausgeführt. Der Abdichtungsanschluss an den Türrahmen wird vor Verdeckung dokumentiert. Der im Vergabeverfahren verworfene vier Zentimeter hohe Standardschwellenanschluss wird nicht eingebaut.',
        ]),
        ('4 Fenster, Türen und Zugänglichkeit', [
            'Die Fenster erhalten eine wärmeschutzgerechte Mehrscheibenverglasung, funktionsfähige Beschläge und die im Plan bezeichneten Öffnungsrichtungen. Als vertraglicher Planungswert für das vollständige Fenster wird ein Wärmedurchgangskoeffizient von höchstens 0,90 W/(m²K) vereinbart. Erforderliche Sicherungen und Lüftungsfunktionen werden mit dem Raum- und Lüftungskonzept abgestimmt. Maße und Teilungen ergeben sich aus den genehmigten und koordinierten Plänen.',
            'Wohnungseingänge erhalten eine lichte Durchgangsbreite von mindestens 0,90 m. Die barrierefreien Wegeketten werden durchgehend ohne unzulässige Schwellen hergestellt. Wohnung 01 erhält zusätzlich die in BF01 eingezeichneten rollstuhlgerechten Bewegungsflächen einschließlich 1,50 m mal 1,50 m im Bad, geeigneter Türanfahrbereiche und unterfahrbaren Waschtischs. Die Waschmaschine ist dort im Abstellbereich vorgesehen, damit die Bewegungsfläche frei bleibt.',
            'Innentüren erhalten eine weiße glatte Oberfläche und beschlägebezogen eine zeitlose Edelstahloptik. Die in der Planung erforderlichen Brandschutz- und Rauchschutzabschlüsse werden mit ihrem nachgewiesenen System eingebaut; dekorative Bemusterung darf diese Eigenschaften nicht verdrängen. Türschließkräfte und nutzbare Breiten sind vor Abnahme zu prüfen.',
        ]),
        ('5 Wärmeschutz und Schallschutz', [
            'Die energetische Planung sieht Außenwand 0,18 W/(m²K), Dach 0,14 W/(m²K), Bodenplatte 0,20 W/(m²K) und Fenster 0,90 W/(m²K) als jeweilige obere Zielwerte vor. Die Verkäuferin schuldet außerdem den vollständigen gesetzlich erforderlichen Nachweis für das Gebäude. Die Einzelwerte allein behaupten weder eine bestimmte Förderklasse noch eine Fördermittelzusage. Eine nachträgliche Förderung setzt ein gesondertes Verfahren voraus.',
            'Als zusätzliche vertragliche Schallschutzziele für die hier betroffenen Wohnungstrennbauteile werden ein bewertetes Bau-Schalldämm-Maß von mindestens 55 dB für Wohnungstrennwände und mindestens 54 dB für Wohnungstrenndecken sowie ein bewerteter Norm-Trittschallpegel von höchstens 46 dB zwischen fremden Wohn- und Schlafräumen vereinbart. Für Geräusche gemeinschaftlicher haustechnischer Anlagen in fremden schutzbedürftigen Räumen wird ein maximaler A-bewerteter Norm-Schalldruckpegel von 27 dB vereinbart, soweit das einschlägige Mess- und Bewertungsverfahren auf die konkrete Anlage anwendbar ist.',
            'Diese Zahlen sind gewählte Vertragsanforderungen der Verkaufsoption, keine Wiedergabe einer ungeprüft behaupteten DIN-Mindesttabelle. Vor Einbeziehung in einen Vertrag muss die Fachplanung ihre Erreichbarkeit mit den tatsächlich gewählten Bauteilen und Anschlüssen nachweisen; erforderliche planerische und kalkulatorische Anpassungen trägt innerhalb des vereinbarten Festpreises die Verkäuferin. Die Prüfmethode und Messorte werden im Schallschutznachweis festgelegt. Die Option darf ohne diesen Abgleich nicht freigegeben werden.',
        ]),
        ('6 Heizung, Warmwasser und Lüftung', [
            'Die gemeinsame Wärmeversorgung erfolgt durch zwei Wärmepumpengeräte mit jeweils 18 kW Nennleistung im abgestimmten Kaskadenbetrieb. Die aus dem bisherigen Berechnungsstand ermittelte Gebäudeheizlast beträgt 27,2 kW. Nennleistung, Auslegungstemperatur und tatsächlicher Betriebspunkt sind voneinander zu unterscheiden; die Geräteauswahl muss den endgültigen fachplanerischen Nachweis erfüllen. Ein Anspruch auf einen bestimmten Jahresstromverbrauch folgt daraus nicht.',
            'Die Wohnungen erhalten die vorgesehenen Flächenheizkreise mit wohnungsbezogener Regelung, dokumentiertem hydraulischem Abgleich und nachvollziehbaren Einstellwerten. Trinkwasserinstallation, Warmwasserbereitung und Lüftung werden hygienegerecht und nach den einschlägigen technischen Regeln geplant, gespült, geprüft und in Betrieb genommen. Die Verkäuferin übergibt die vollständigen Prüf- und Einstellunterlagen. Gemeinschaftliche Wartung und individuelle Bedienpflichten sind in den Betriebsunterlagen getrennt beschrieben.',
        ]),
        ('7 Sanitärausstattung und Oberflächen', [
            'Jede Wohnung erhält ein fertig nutzbares Bad mit bodengleichem Duschbereich, WC, Waschtisch, Armaturen und den geplanten Anschlüssen. Die Basisausstattung verwendet weiße Keramik, verchromte beziehungsweise gleichwertig dauerhaft geschützte Armaturen und eine neutral helle Fliesenserie. Im rollstuhlgerechten Bad der Wohnung 01 werden die erforderlichen Bewegungsflächen, Befestigungsuntergründe für Haltegriffe und die unterfahrbare Anordnung eingehalten.',
            'Wohn- und Schlafräume erhalten den in der Bemusterungsordnung bezeichneten Eiche-Natur-Mehrschichtboden, Flure eine darauf abgestimmte strapazierfähige Oberfläche, Nassbereiche den vorgesehenen keramischen Belag. Sockelleisten, Übergänge, dauerelastische Anschlussfugen und die erforderlichen Untergrundvorbereitungen gehören zur Leistung. Innenwände erhalten eine geglättete, streichfähige Oberfläche und einen deckenden weißen Schlussanstrich. Spachtelqualität und Untergrund werden auf diese Ausführung abgestimmt; besondere Streiflichtperfektion ist nur bei gesonderter Vereinbarung geschuldet.',
            'Die Käuferseite kann innerhalb der rechtzeitig angebotenen, kostenneutralen Basisgruppen wählen. Es werden keine verdeckten Quadratmeterzulagen für die vereinbarte Grundausführung verlangt. Abweichende Formate, Diagonalverlegung, Nischen, besondere Untergründe oder ein Produkt außerhalb der Basisgruppe werden erst nach vollständigem Sonderwunschangebot ausgeführt.',
        ]),
        ('8 Elektro, Kommunikation, PV und Laden', [
            'Die Wohnungen erhalten getrennte Unterverteilungen, die nach dem abgestimmten Elektroplan erforderlichen Steckdosen, Schalter und Lichtauslässe sowie Schutzorgane und Potentialausgleich. Raumbezogene Stromkreise und Zählerzuordnung werden eindeutig beschriftet. Küchenanschlüsse werden nach dem vorhandenen Installationsplan vorgerichtet; die konkrete Einbauküche ist nicht enthalten. Die Fertigstellungsprüfung umfasst alle neu errichteten Anlagenbereiche und wird mit Messwerten dokumentiert.',
            'Auf den Dächern werden 80 Module mit je 450 Wp, insgesamt 36 kWp und 160 m² Modulfläche, samt Wechselrichter, Verkabelung und Betriebsunterlagen installiert. Das Anlagenkonzept ist auf Allgemeinstrom und Überschusseinspeisung ausgerichtet. Die steuerliche Behandlung der PV-Lieferung und die spätere Betreiberbesteuerung sind getrennte Fragen; die Kaufpreisvereinbarung enthält keine Garantie eines dauerhaften Nullsteuersatzes für sämtliche Leistungen.',
            'Alle acht Stellplätze erhalten die vereinbarte Leitungsinfrastruktur und Vorverkabelung. Zwei Ladepunkte mit jeweils bis zu 11 kW werden einschließlich Lastmanagement installiert, davon einer am barrierefreien Stellplatz 01. Für die Verkaufsoption ist der zweite Ladepunkt dem Stellplatz 02 zugeordnet. Diese zusätzliche konkrete Zuordnung ist vor Vertragsfreigabe im Elektro- und Aufteilungsplan zu übernehmen. Der tatsächlich gleichzeitig verfügbare Ladestrom hängt von Netzanschluss und Lastmanagement ab.',
        ]),
        ('9 Aufzug, Außenanlagen und Übergabe', [
            'Der Aufzug erschließt drei Haltestellen und hat eine vereinbarte nutzbare Kabine von 1,10 m mal 1,40 m sowie eine lichte Türbreite von 0,90 m bei 630 kg Nennlast. Notruf, Prüfungen und Betriebsfreigabe gehören zum herzustellenden funktionsfähigen System. Die Zugänge dürfen nicht durch Türschwenkbereiche oder Abstellnutzung eingeschränkt werden.',
            'Die Außenanlagen umfassen acht nutzbare Pkw-Stellplätze, sichere Wege, 16 gesicherte Bewohner-Fahrradplätze, vier Besucherplätze, einen 30 m² großen Spielbereich und die planmäßig vorgesehenen Grünflächen. Die Verkäuferin übergibt die Anlagen in einem nutzbaren Zustand; Verkehrs- und Spielplatzsicherheit dürfen nicht auf eine unbestimmte spätere Ausführung verschoben werden. Anwuchs- und Pflegeleistungen folgen dem vereinbarten Außenanlagenauftrag.',
            'Zur Übergabe werden Schlüssel, Zählerstände, Anlagenbestände und Betriebsunterlagen in einer wohnungsbezogenen Liste festgehalten. Die Beseitigung protokollierter Mängel erfolgt mit nachvollziehbaren Erledigungsnachweisen. Der Betrieb einer Anlage ersetzt nicht die Erklärung der Abnahme. Wartungs- und Prüftermine werden der Verwaltung mitgeteilt; ein nicht bestehender Wartungsvertrag wird nicht durch einen Anlagenordner fingiert.',
        ]),
    ]
    for title, paras in sections:
        h(d, title)
        for text in paras: p(d,text)
    h(d, '10 Planverzeichnis und Einbeziehung')
    table(d, [['Plan', 'Gegenstand', 'Bezugsstand'],
        ['LP01 Revision 1', 'Lageplan und Außenflächen', '24.06.2027'],
        ['GP01', 'Grundrisse Erdgeschoss und erstes Obergeschoss', '19.03.2027, Türpräzisierungen 24.06.2027'],
        ['GP02', 'Grundriss zweites Obergeschoss', '19.03.2027, Türpräzisierungen 24.06.2027'],
        ['SP01', 'Schnitt und Ansichten', '19.03.2027'], ['BF01 Revision 1', 'Barrierefreiheit Wohnung 01 und Aufzug', '24.06.2027'],
        ['TE01', 'Dach und Entwässerung', '24.06.2027'], ['AP02 und AP03', 'Schwelle und Deckendurchbruch', '21.08.2027'],
        ['AP-OPTION-01', 'Noch zu bescheinigender Aufteilungsplan mit Stellplatzzuordnung', 'Vorbereitung 01.11.2027']], [3.2,6.6,4.2])
    p(d, 'Die tatsächlich beigefügten Pläne müssen genau diesen Ständen entsprechen. '
         'Für die endgültige Verkaufsurkunde sind die vorbereitete Aufteilung, '
         'die technisch geprüften Zusatzanforderungen und die unverändert '
         'geschuldeten Merkmale gemeinsam freizugeben. Die Verkäuferin darf '
         'durch spätere Planrevisionen keine schon vereinbarte Beschaffenheit '
         'ohne wirksame Änderungsgrundlage verdrängen.')
    save(d, '02_Gemeinsame_Anlagen', 'AN02_Baubeschreibung_BT_BB_01.docx')


def annexes():
    for u in UNITS:
        n=u['n'];code=f'WE{n:02d}'
        d=start(f'Raum- und Zuordnungsliste · {code}')
        p(d, f'Anlage zum Bauträgervertragsentwurf für {u["names"]}. '
             'Diese Wohnung ist Bestandteil der noch nicht beschlossenen Verkaufsoption; '
             'sie wird dadurch nicht aus dem Hauptprojekt vermietet oder tatsächlich verkauft.')
        h(d,'1 Räume und Flächen')
        table(d,[['Raum','Wohnfläche m²']]+[[room,euro(area)] for room,area in rooms(u)]+[['Summe',euro(u['area'])]],[10,4])
        p(d,f'Die Wohnung liegt im {u["floor"].replace("es Obergeschoss", "en Obergeschoss")}, {u["side"]}. '
            'Alle Räume liegen innerhalb der Wohnung; ein Kellerraum und ein '
            'Dachterrassenanteil sind nicht enthalten. Die Räume werden vollständig '
            'in der Wohnflächenberechnung angesetzt. Der Aufteilungsplan muss '
            'die genaue Abgrenzung und die nummerngleiche Zuordnung erkennen lassen.')
        h(d,'2 Rechtliche Zuordnung')
        p(d,f'Miteigentumsanteil: {u["area"]}/600. Zugeordnetes Sondernutzungsrecht: '
            f'Stellplatz {n:02d}. Die Außen- und Technikflächen bleiben '
            'im Übrigen gemeinschaftlich. Gemeinschaftliche Bauteile innerhalb '
            'der Wohnung bleiben ungeachtet ihrer räumlichen Lage Gemeinschaftseigentum.')
        p(d,('Wohnung 01 und die zugehörige Wegekette sind zusätzlich rollstuhlgerecht '
             'auszuführen; Stellplatz 01 ist der gekennzeichnete barrierefreie Stellplatz. '
             'Die Waschmaschine wird im Abstellraum angeschlossen, damit das Bad '
             'die vereinbarten freien Bewegungsflächen behält.' if n==1 else
             'Die Wohnung wird barrierefrei hergestellt. Die nur für Wohnung 01 '
             'zusätzlich vereinbarte rollstuhlgerechte Ausstattung ist für diese '
             'Wohnung nicht pauschal mitverkauft; die nutzbaren Türbreiten und '
             'Wegeketten aus dem gemeinsamen Konzept bleiben geschuldet.'))
        h(d,'3 Kaufpreis und Ausstattung')
        p(d,f'Der wohnungsbezogene Gesamtpreis einschließlich Stellplatzrecht beträgt '
            f'{euro(u["price"])} EUR. Enthaltene Grundausstattung und '
            'Auswahlmöglichkeiten folgen BT-BB-01 und BT-BM-01. '
            'Ein kostenpflichtiger Sonderwunsch ist für diese Wohnung '
            'noch nicht beauftragt. Vor Vertragsfreigabe sind '
            'Käuferdaten, Wohnungslage, Fläche, Planstand und '
            'Sondernutzungsnummer gemeinsam abzugleichen.')
        save(d,'03_Wohnungsanlagen',f'WE{n:02d}_Raumliste_und_Zuordnung.docx')
        d=start(f'Zahlungs- und Sicherungsplan · {code}')
        p(d,'Dieser Plan ist eine Vertragsanlage und keine Rechnung, '
            'Fälligkeitsmitteilung oder Zahlungsfreigabe. Solange die '
            'Verkaufsoption nicht wirksam aktiviert und beurkundet ist, '
            'entsteht daraus kein Zahlungsanspruch.')
        table(d,rate_rows(u),[1.4,8.2,1.6,3.2])
        h(d,'1 Berechnung der Bündel')
        for number,description,percent,calc in RATES: p(d,f'Rate {number}: {calc}')
        h(d,'2 Fertigstellungssicherheit')
        p(d,f'Fünf Prozent Sicherheit entsprechen {euro(Decimal(u["price"])*Decimal(".05"))} EUR. '
            f'Bei Einbehalt beträgt die erste Auszahlung statt '
            f'{euro(Decimal(u["price"])*Decimal(".30"))} EUR höchstens '
            f'{euro(Decimal(u["price"])*Decimal(".25"))} EUR. '
            'Der Einbehalt wird nicht ohne Wegfall seines Sicherungszwecks freigegeben. '
            'Ein geeigneter Garantieersatz, Mängeleinbehalte und die '
            'allgemeinen MaBV-Sicherungen sind getrennt zu prüfen.')
        save(d,'03_Wohnungsanlagen',f'WE{n:02d}_Raten_und_Sicherungsplan.docx')


def management_documents():
    d=start('Beschlussvorlage · Prüfung der Einzelverkaufsoption')
    h(d,'1 Gegenstand der Entscheidung')
    p(d,'Die Geschäftsführung legt den Gesellschaftern am 1. November 2027 '
        'die Möglichkeit vor, das geplante Mietobjekt vor Baubeginn '
        'in acht Wohnungseigentumsrechte aufzuteilen und einzeln '
        'zu veräußern. Das bisher freigegebene Vermietungskonzept '
        'gilt weiter. Die nachfolgenden Erwerberprofile und Preise '
        'dienen der vorbereiteten Vertragsgestaltung; es liegen '
        'weder verbindliche Angebote noch Reservierungen oder Kaufverträge vor.')
    h(d,'2 Vorgesehene wirtschaftliche Eckdaten')
    table(d,[['Einheit','Ansatz Gesamtpreis EUR','Käuferprofil für Entwurf']]+[
        [f'WE{u["n"]:02d}',euro(u['price']),u['names']] for u in UNITS]+
        [['Gesamt',euro(sum(u['price'] for u in UNITS)),'Acht voneinander getrennte Erwerbsvorgänge']], [2,4.2,7.8])
    p(d,'Die Ansätze sind fiktive Verhandlungspreise und keine Marktwertermittlung. '
        'Aus ihrem Vergleich mit dem bisherigen Mietprojektbudget ergibt sich '
        'noch kein belastbarer Veräußerungsgewinn. Vertrieb, Teilung, '
        'Abgeschlossenheit, mögliche Zusatzkosten der Schallschutzanforderungen, '
        'Steuern, Finanzierung, Pfandfreigabe und Ausfallrisiken sind '
        'vor einer Freigabe gesondert zu kalkulieren. Die Projektbank '
        'hat noch keine auf diese Verkaufsoption bezogene Freistellung abgegeben.')
    h(d,'3 Voraussetzungen eines späteren Beschlusses')
    p(d,'Vor einer endgültigen Entscheidung sind eine marktbezogene '
        'Preisprüfung, eine aktualisierte Liquiditätsplanung ohne '
        'Doppelzählung von Mieten und Verkaufserlösen, die schriftliche '
        'Abstimmung mit der Projektbank und die rechtlich-technische '
        'Teilungsprüfung vorzulegen. Die notarielle Entwurfsfassung '
        'ist mit dem tatsächlichen Käufer und dem zum Abschluss '
        'geltenden Recht zu prüfen. Eine Vertriebsfreigabe '
        'darf nicht als Baufreigabe oder Kaufpreisfälligkeit behandelt werden.')
    h(d,'4 Zur Entscheidung gestellter Beschlusstext')
    p(d,'Die Gesellschafterversammlung beschließt, die Einzelverkaufsoption '
        'erst nach Vorlage und Zustimmung zu den in Ziffer 3 '
        'bezeichneten Unterlagen gesondert freizugeben. Bis '
        'dahin bleibt es bei der bestehenden Vermietungsplanung. '
        'Die Geschäftsführung darf in der Zwischenzeit unverbindliche '
        'Entwürfe und Kalkulationen vorbereiten, jedoch keine '
        'Kaufpreiszahlungen einziehen oder einen bereits '
        'beschlossenen Verkauf mitteilen.')
    p(d,'Abstimmung und Unterschriften sind offen. Diese Vorlage '
        'dokumentiert noch keinen gefassten Beschluss.')
    save(d,'00_Entscheidungsvorbereitung','OP01_Beschlussvorlage_Einzelverkauf.docx')
    d=start('Vorbereitung des Aufteilungsplans · AP-OPTION-01')
    h(d,'1 Planauftrag')
    p(d,'Die bestehende Objektplanung soll für den Fall einer späteren '
        'Freigabe der Verkaufsoption einen eindeutig nummerierten '
        'Aufteilungsplan vorbereiten. Grundlage sind die bestehenden '
        'Grundrisse GP01 und GP02, der Lageplan LP01 Revision 1 '
        'und die genehmigte Nutzung. Eine Änderung der Kubatur '
        'oder der acht Wohnungszuschnitte wird damit nicht beauftragt.')
    h(d,'2 Darzustellende Rechte und Flächen')
    p(d,'Jeder zu einer Wohnung gehörende Raum muss dieselbe Nummer '
        'von 01 bis 08 tragen. Gemeinschaftliche Verkehrs- und '
        'Technikflächen sind eindeutig abzugrenzen. Im Lageplan '
        'sind die acht Stellplatzsonder­nutzungsflächen nummeriert '
        'und vermaßt darzustellen. Das barrierefreie Stellplatzrecht '
        '01 ist der Wohnung 01 zugeordnet; Stellplatz 02 erhält '
        'im Optionsstand den zweiten Ladepunkt. Fahrrad-, '
        'Spiel-, Wartungs- und Rettungsflächen bleiben gemeinschaftlich.')
    h(d,'3 Behördliches Verfahren und Vollzug')
    p(d,'Aufteilungsplan und Abgeschlossenheitsbescheinigung werden '
        'im gesetzlich vorgesehenen Verfahren beantragt. Eine '
        'bloße Grundrisskopie ohne die notwendigen '
        'behördlichen Merkmale ist nicht als bereits '
        'bescheinigter Plan zu verwenden. Eingangsbestätigung, '
        'Nachforderung, Bescheinigung und notarieller '
        'Einreichungsstand werden erst dokumentiert, '
        'wenn sie tatsächlich im Optionsverfahren vorliegen.')
    save(d,'00_Entscheidungsvorbereitung','OP02_Aufteilungsplan_Vorbereitungsauftrag.docx')
    d=start('Anfrageentwurf an die Projektbank · Einzelpfandfreistellung')
    p(d,'An die Leinebogen Projektbank AG, Firmenkundenbetreuung. '
        'Betreff: SW-HI-26-08, unverbindliche Prüfung einer '
        'Einzelverkaufsoption, Projektgrundschuld 1.600.000,00 EUR.')
    p(d,'Sehr geehrte Frau West,')
    p(d,'wir prüfen derzeit als noch nicht beschlossene Alternative '
        'zur Vermietung den Einzelverkauf der acht Wohnungen '
        'vor Baubeginn. Bitte teilen Sie uns mit, unter '
        'welchen wirtschaftlichen Voraussetzungen Sie '
        'für jede Einheit eine MaBV-konforme Freistellung '
        'der nicht zu übernehmenden Grundpfandrechte '
        'auch für den Fall unvollendeter Bauausführung '
        'erklären könnten. Wir bitten um einen '
        'vollständigen Wortlaut für die notarielle Prüfung.')
    p(d,'Bitte trennen Sie die projektinterne Darlehensrückführung '
        'von dem gegenüber dem jeweiligen Erwerber '
        'auszuhändigenden Freistellungsversprechen. '
        'Benennen Sie die vorgesehenen Ablösebeträge, '
        'Konten, Rangfolgen und eine etwaige '
        'Rückzahlungsalternative im gesetzlich '
        'zulässigen Umfang. Die beigefügte Preisübersicht '
        'enthält lediglich Planansätze, keine '
        'beurkundeten Kaufpreise.')
    p(d,'Bis zur gesonderten Entscheidung werden keine '
        'Erwerberzahlungen eingezogen und keine '
        'bestehenden Darlehenspflichten geändert. '
        'Bitte bestätigen Sie gesondert, welche '
        'Zustimmungen aus dem Projektfinanzierungsvertrag '
        'vor einer etwaigen Veräußerung benötigt werden.')
    p(d,'Mit freundlichen Grüßen\nMaren Birk\nSteinbogen Wohnen GmbH')
    p(d,'Versandstatus: Entwurf, noch nicht versandt.')
    save(d,'00_Entscheidungsvorbereitung','OP03_Bankanfrage_Pfandfreistellung_Entwurf.docx')


def bemusterung():
    d=start('Bemusterungs- und Sonderwunschordnung · BT-BM-01')
    h(d,'1 Grundausstattung und Auswahlfenster')
    p(d,'Die Grundausstattung besteht aus weißer Sanitärkeramik, '
        'neutral hellen rechteckigen Fliesen im Bad, '
        'einem Eiche-Natur-Mehrschichtboden in Wohn- '
        'und Schlafräumen, weißen glatten Innentüren '
        'und einer weißen Schalterserie. Die '
        'Verkäuferin stellt vor Auswahl ein konkretes '
        'Bemusterungsblatt mit Hersteller, Produkt, '
        'Musterkennzeichen und technischen Eigenschaften '
        'zur Verfügung. Die Eigenschaften der '
        'vereinbarten Baubeschreibung bleiben verbindlich.')
    p(d,'Die Käuferseite erhält für eine angebotene '
        'kostenneutrale Auswahl mindestens vierzehn '
        'Kalendertage. Eine Erinnerung bezeichnet '
        'die noch ausstehende Entscheidung und '
        'räumt eine angemessene Nachfrist ein. '
        'Bleibt die Auswahl danach aus, wird '
        'die bereits konkret beschriebene '
        'Grundleistung ausgeführt. Schweigen '
        'beauftragt keinen kostenpflichtigen Zusatz.')
    h(d,'2 Sonderwunschangebot')
    p(d,'Das Sonderwunschangebot enthält die betreffende '
        'Wohnung, den Raum, die ursprüngliche '
        'Leistung, die gewünschte Ersatz- oder '
        'Zusatzleistung, den Bruttopreis, ersparte '
        'Grundleistungen, Auswirkungen auf '
        'technische Nachweise und den konkreten '
        'Terminplan. Die Verkäuferin weist '
        'einen nicht sicher kalkulierbaren '
        'Wunsch zunächst zur Klärung zurück, '
        'statt ihn unbestimmt anzunehmen.')
    p(d,'Die Käuferseite beauftragt erst auf Grundlage '
        'des vollständigen Angebots. Die '
        'gesetzlich erforderliche Form ist '
        'vor Ausführung mit dem Notariat '
        'zu prüfen. Planungs- oder '
        'Prüfaufwand ohne spätere '
        'Ausführung ist nur vergütungspflichtig, '
        'wenn dies zuvor gesondert '
        'vereinbart wurde. Zahlungen '
        'dürfen die MaBV-Sicherungen '
        'nicht durch eine Nebenrechnung umgehen.')
    h(d,'3 Technische Grenzen und Freigabe')
    p(d,'Änderungen an tragenden Bauteilen, '
        'gemeinschaftlichen Anlagen, Rettungswegen, '
        'Abdichtungen, barrierefreien Flächen '
        'und Fassaden werden nur nach '
        'fachlicher Prüfung und den '
        'erforderlichen Genehmigungen '
        'ausgeführt. Eine Unterzeichnung '
        'des Bemusterungsblatts ersetzt '
        'diese Prüfungen nicht. '
        'Verschobene Leistungen werden '
        'im Planregister und in '
        'der Schnittstellenliste '
        'nachvollziehbar fortgeführt.')
    h(d,'4 Feststellung bei Übergabe')
    p(d,'Die gewählte Ausstattung wird wohnungsbezogen '
        'mit Produktbezeichnung und '
        'Freigabestand in die '
        'Übergabeunterlagen übernommen. '
        'Eine Materialauswahl erklärt '
        'keine Abnahme der späteren '
        'Einbauleistung. Reklamationen '
        'werden anhand des tatsächlich '
        'vereinbarten Produkts und '
        'der Ausführung geprüft.')
    save(d,'02_Gemeinsame_Anlagen','AN03_Bemusterung_Sonderwunschordnung.docx')


def templates():
    specs = [
        ('V02_Abnahmeprotokoll_Sondereigentum.docx','Abnahmeprotokoll Sondereigentum', [
            ('1 Beteiligte und Gegenstand','Am [Datum] besichtigen die Verkäuferin [Name], die Käuferseite [Name] und die hinzugezogene Vertrauensperson [Name] die Wohnung [Nummer]. Grundlage sind der Erwerbsvertrag vom [Datum], der Planstand [Bezeichnung] und die vereinbarten Sonderwünsche [Bezeichnung].'),
            ('2 Feststellungen','Die Beteiligten prüfen die zugänglichen Räume, Oberflächen, Türen, Fenster, Installationen und die wohnungsbezogenen Betriebsfunktionen. Nicht zugängliche oder nicht geprüfte Bereiche werden mit Grund und Nachholtermin bezeichnet. Die beigefügte Mängelliste enthält für jeden Befund Raum, Bauteil, Erscheinung, Einordnung, Beleg und vorgesehenen Erledigungstermin.'),
            ('3 Individuelle Erklärung','Die Käuferseite erklärt nach Prüfung: [Abnahme mit konkret bezeichneten Vorbehalten oder begründete Verweigerung wegen wesentlicher Mängel]. Rechte wegen der in der Anlage als bekannt bezeichneten Mängel werden ausdrücklich vorbehalten. Ein etwaiger Vorbehalt wegen einer gesondert vereinbarten Vertragsstrafe wird eigenständig erklärt. Eine Zustimmung der Verkäuferin zu jedem Befund ist nicht Voraussetzung für dessen Dokumentation.'),
            ('4 Nacharbeiten und Unterlagen','Die Verkäuferin stimmt den Zugang für die bezeichneten Nacharbeiten mit der Käuferseite ab. Die technische Erledigung wird nachgeprüft. Dieses Protokoll ersetzt weder die Abnahme des Gemeinschaftseigentums noch die Prüfung der Schlussratenfälligkeit. Übergebene Unterlagen und Schlüssel werden getrennt im Übergabeprotokoll aufgeführt.'),
        ]),
        ('V03_Abnahmeprotokoll_Gemeinschaftseigentum.docx','Individuelle Abnahmeerklärung Gemeinschaftseigentum', [
            ('1 Begehung und Vertragsbezug','Am [Datum] findet die Begehung des Gemeinschaftseigentums zum Erwerbsvertrag der Käuferseite [Name], Wohnung [Nummer], statt. Der Gegenstand umfasst [konkret bezeichnete Bereiche]. Der technische Sachverständige [Name] begleitet die Prüfung, erklärt dadurch aber keine Abnahme im Namen der Käuferseite.'),
            ('2 Prüfungsumfang','Dach, Fassade, tragende und gemeinschaftliche Bauteile, Technik, Aufzug, Außenanlagen und die erforderlichen Betriebsunterlagen werden anhand der beigefügten Liste geprüft. Nicht untersuchte oder noch unzugängliche Teile sind ausdrücklich aufgeführt. Eine frühere Erklärung anderer Erwerber wird nicht als Erklärung dieser Käuferseite übernommen.'),
            ('3 Erklärung der Käuferseite','Die Käuferseite erklärt persönlich oder durch eine gesondert und frei bevollmächtigte Vertrauensperson: [konkrete Abnahmeentscheidung]. Die in der Befundliste bezeichneten bekannten Mängel und Restleistungen werden vorbehalten. Eine Mehrheit der Eigentümer und die Verwaltung ersetzen diese Erklärung nicht.'),
            ('4 Folgeverfahren','Die Verkäuferin teilt für jeden bestätigten Mangel den vorgesehenen Erledigungstermin mit. Streitig gebliebene Befunde und die hierzu vorhandenen Unterlagen werden ohne Anerkenntniszwang festgehalten. Eine technische Nachkontrolle erklärt weder einen pauschalen Verzicht noch einen gemeinsamen Verjährungsbeginn für sämtliche Erwerber.'),
        ]),
        ('V04_Sonderwunschvereinbarung.docx','Sonderwunschvereinbarung mit Preis- und Terminfolge', [
            ('1 Vertragsbezug und gewünschte Änderung','Die Parteien des Erwerbsvertrags vom [Datum] über Wohnung [Nummer] vereinbaren nach Prüfung der erforderlichen Form folgende Änderung: [vollständige Beschreibung der neuen Leistung]. Sie ersetzt die bisher vereinbarte Leistung [Beschreibung] im bezeichneten Umfang.'),
            ('2 Vergütungsfolge','Der Bruttopreis der neuen Leistung beträgt [Betrag] EUR. Für entfallende Grundleistungen werden [Betrag] EUR abgezogen. Die zusätzliche vereinbarte Vergütung beträgt daher [Betrag] EUR. Diese Vereinbarung schafft keine Zahlungspflicht vor Eintritt der gesetzlichen und vertraglichen Fälligkeitsvoraussetzungen.'),
            ('3 Termine und Freigaben','Die konkret nachgewiesene Auswirkung auf den kritischen Bauablauf beträgt [Dauer und betroffene Arbeiten]. Als neuer verbindlicher Termin für die betroffene Leistung wird [Datum] vereinbart; andere Termine bleiben unberührt, soweit sie nicht ausdrücklich bezeichnet werden. Die erforderlichen technischen und behördlichen Freigaben [Bezeichnung] sind vor Ausführung einzuholen.'),
            ('4 Dokumentation','Die Änderung wird im Planregister und im wohnungsbezogenen Ausstattungsblatt fortgeschrieben. Die Parteien erhalten den vollständigen Text mit Anlagen. Die Ausführung beginnt erst nach wirksamer Vereinbarung und Klärung der gesetzlichen Form. Alle nicht geänderten Bestimmungen des Erwerbsvertrags gelten fort.'),
        ]),
        ('V05_MaBV_Ratenabruf.docx','Ratenabruf nach dokumentiertem Bautenstand', [
            ('1 Forderung und Belege','Aus dem wirksamen Erwerbsvertrag vom [Datum] über Wohnung [Nummer] wird die Rate [Nummer] in Höhe von [Betrag] EUR angefordert. Die Rate beruht auf den vollständig erreichten Leistungen [konkrete Leistungen]. Der beigefügte Bericht [Datum und Verfasser] beschreibt den jeweiligen Fertigstellungsstand.'),
            ('2 Allgemeine Voraussetzungen','Die notarielle Mitteilung vom [Datum], die Eintragung des Wohnungseigentums und der Vormerkung sowie die ausgehändigte Freistellungserklärung vom [Datum] sind zugeordnet. Die notwendige Baugenehmigung und ihre für die Ausführung maßgeblichen Voraussetzungen liegen vor. Fehlt eine dieser Voraussetzungen, darf dieses Schreiben nicht als Zahlungsaufforderung freigegeben werden.'),
            ('3 Sicherheit und Zahlung','Von dem Ratenbetrag sind die noch aufzubauende Fertigstellungssicherheit von [Betrag] EUR und begründete sonstige Abzüge gesondert abzusetzen. Der danach angeforderte Betrag beträgt [Betrag] EUR. Die Zahlung wird frühestens nach Ablauf der vertraglichen Frist und nur auf die verifizierte, nach der notariellen Abwicklung zulässige Bankverbindung verlangt. Gesetzliche Einwendungen und Zurückbehaltungsrechte bleiben unberührt.'),
            ('4 Schlussrate','Wird die Schlussrate angefordert, sind die vollständige Fertigstellung, die Erledigung protokollierter Mängel und Restarbeiten sowie die noch offenen streitigen Befunde konkret zu prüfen. Die bloße Erklärung der Abnahme ist dafür kein Ersatz.'),
        ]),
        ('V06_Maengelanzeige_und_Terminabstimmung.docx','Mängelanzeige und Terminabstimmung', [
            ('1 Befund','In Wohnung beziehungsweise Gemeinschaftsbereich [Bezeichnung] wurde am [Datum] folgender Befund festgestellt: [konkrete Erscheinung]. Betroffen sind [Bauteil und Ort]. Die beigefügten Unterlagen [Bezeichnung] dokumentieren den Zustand; eine technische Ursache wird nur insoweit benannt, wie sie tatsächlich festgestellt ist.'),
            ('2 Abhilfeverlangen','Die Verkäuferin wird gebeten, den Befund zu prüfen und den geschuldeten vertragsgemäßen Zustand bis zum [angemessene konkrete Frist] herzustellen. Für eine gemeinsame Besichtigung werden [Termine] angeboten. Sofern die vorgeschlagene Frist aus nachweisbaren technischen Gründen nicht genügt, wird unverzüglich um einen begründeten Ablaufvorschlag gebeten.'),
            ('3 Dringlichkeit','Bei einer Gefahr für Personen oder drohender Ausweitung des Schadens sind geeignete Sofortmaßnahmen zu veranlassen und zu dokumentieren. Solche Maßnahmen erklären kein Anerkenntnis einer bestimmten Kostenverteilung. Erforderliche Beweise werden vor Veränderung des Zustands gesichert, soweit die Gefahrenabwehr dies zulässt.'),
            ('4 Rechte und Kommunikation','Dieses Schreiben enthält keinen Verzicht auf weitergehende gesetzliche Rechte. Die Antwort soll Befund, Verantwortlichkeit, vorgesehenen Arbeitsschritt und Termin getrennt behandeln. Die Mängelliste wird nach tatsächlicher Erledigung und Nachkontrolle fortgeschrieben.'),
        ]),
    ]
    for filename,title,sections in specs:
        d=new_document(title,STAND)
        p(d,'Ausfüllbare Arbeitsvorlage zum gesonderten Bauträgererwerb.')
        for heading,text in sections:h(d,heading);p(d,text)
        p(d,'Unterschrift beziehungsweise dokumentierte Erklärung: ______________________________')
        save_docx(d,Path('12_Wordvorlagen/Bautraeger')/filename)


def write_register():
    out=CASE/OPTION/'00_Entscheidungsvorbereitung'/'OP04_Preisansatz_und_Ratenkontrolle.csv'
    stream=io.StringIO(newline='')
    writer=csv.writer(stream,delimiter=';')
    writer.writerow(['Status','Einheit','Wohnfläche_m2','MEA_Zähler','MEA_Nenner','Planpreis_EUR',
        'Sicherheit_5Prozent_EUR','Erste_Auszahlung_25Prozent_EUR']+[f'Rate{i}_EUR' for i in range(1,8)])
    for u in UNITS:
        price=Decimal(u['price'])
        writer.writerow(['Option_nicht_beschlossen',f'WE{u["n"]:02d}',u['area'],u['area'],600,
            euro(price),euro(price*Decimal('.05')),euro(price*Decimal('.25'))]+[
            euro(price*r[2]/100) for r in RATES])
    out.parent.mkdir(parents=True,exist_ok=True);out.write_text(stream.getvalue(),encoding='utf-8-sig')
    QUALITY.mkdir(parents=True,exist_ok=True)
    (QUALITY/'verkaufsoption-daten.json').write_text(json.dumps({
        'status':'nicht beschlossen; keine Beurkundung, Zahlung oder Buchung',
        'optionsstichtag':STAND,'rechtsstand':'2026-09-28','units':UNITS,
        'rate_percent':[str(r[2]) for r in RATES],
        'prices_total':sum(u['price'] for u in UNITS),
        'mea_total':'600/600','main_scenario':'Vermietung unverändert',
    },ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def build():
    assert sum(u['area'] for u in UNITS)==600
    assert sum(r[2] for r in RATES)==100
    management_documents();teilung();baubeschreibung();bemusterung();annexes()
    for unit in UNITS:contract(unit)
    contract(UNITS[0],master=True);templates();write_register()
    print('Bauträgeroption: acht Vertragsentwürfe, gemeinsame Anlagen, 16 Wohnungsanlagen, sechs Wordvorlagen; keine Verkaufsbuchung.')


if __name__=='__main__':build()
