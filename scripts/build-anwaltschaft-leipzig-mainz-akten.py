#!/usr/bin/env python3
"""Erzeugt die beiden Arbeitsakten für 445.1.0 ausschließlich in ihren Aktenordnern.

Aufruf: python3 scripts/build-anwaltschaft-leipzig-mainz-akten.py
Danach: node scripts/build-anwaltschaft-leipzig-mainz-workbooks.mjs
Option --pakete baut mit den bestehenden Exportern die Lesefassungen und lokale
Archivkopien unter /tmp. --qa rendert Einzelunterlagen dort zur Sichtkontrolle.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import os
import re
import sys
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from akten_docx_format import separate_section_headings
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether, PageBreak

ROOT = Path(__file__).resolve().parents[1]
QA_ROOT = Path(os.environ.get("AKTEN_QA_DIR", "/tmp/anwaltschaft-leipzig-mainz-445-1-0"))
SLUGS = ("anwaltschaft-lieferstreit-kaffeeroesterei-leipzig", "anwaltschaft-arbeitsrecht-vertrieb-mainz")
NOTICE = "Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\nThis test case file was generated with AI and is an experiment. Use at your own responsibility and risk."
ELSTER = "Elsterbogen Rösterei GmbH\nWeißenfelser Straße 42, 04229 Leipzig\nGeschäftsführerin Friederike Ott | Telefon 0341 247 61 20\nAG Leipzig HRB 39816 | kontakt@elsterbogen-roesterei.de"
NORD = "Nordwerk Verpackungstechnik GmbH\nBillbrookdeich 168, 22113 Hamburg\nGeschäftsführer Hendrik Behn | Telefon 040 731 68 40\nAG Hamburg HRB 162840 | service@nordwerk-verpackung.de"
MUHL = "Mühlentor Sensortechnik GmbH\nRheinallee 88, 55120 Mainz\nGeschäftsführerin Marlene Krämer | Telefon 06131 487 62 0\nAG Mainz HRB 49126 | personal@muehlentor-sensor.de"
HESSE = "Torben Hesse\nWallaustraße 37\n55118 Mainz"
CONTACTS = {
    "ott": ("Friederike Ott", "friederike.ott@elsterbogen-roesterei.de", "Geschäftsführerin", ELSTER),
    "brandt": ("Maja Brandt", "maja.brandt@elsterbogen-roesterei.de", "Produktion", ELSTER),
    "pohl": ("Leonie Pohl", "leonie.pohl@elsterbogen-roesterei.de", "Buchhaltung", ELSTER),
    "behn": ("Hendrik Behn", "hendrik.behn@nordwerk-verpackung.de", "Geschäftsführer", NORD),
    "haas": ("Steffen Haas", "steffen.haas@nordwerk-verpackung.de", "Serviceleitung", NORD),
    "stahl": ("Dr. Hannah Stahl", "hannah.stahl@kanzlei-stahl-leipzig.de", "Rechtsanwältin", "Kanzlei Stahl\nKarl-Heine-Straße 54, 04229 Leipzig"),
    "kraemer": ("Marlene Krämer", "marlene.kraemer@muehlentor-sensor.de", "Geschäftsführerin", MUHL),
    "berg": ("Nils Berg", "nils.berg@muehlentor-sensor.de", "Vertriebsleitung", MUHL),
    "engel": ("Derya Engel", "derya.engel@muehlentor-sensor.de", "Personal und Finanzen", MUHL),
    "hesse": ("Torben Hesse", "torben.hesse@hesse-mainz-post.de", "", HESSE),
    "hesse_work": ("Torben Hesse", "torben.hesse@muehlentor-sensor.de", "Vertrieb", MUHL),
    "reuter": ("Sabine Reuter", "sabine.reuter@reuter-arbeitsrecht-mainz.de", "Rechtsanwältin", "Kanzlei Reuter\nKaiserstraße 51, 55116 Mainz"),
    "folie": ("Ute Schenk", "ute.schenk@saaleflex.de", "Anwendungstechnik", "Saaleflex Folien GmbH\nDieselstraße 19, 06112 Halle (Saale)\nTelefon 0345 681 27 14"),
    "kunde1": ("Jana Seidel", "einkauf@kantine-am-bogen.de", "Einkauf", "Kantine am Bogen GmbH\nNonnenstraße 14, 04229 Leipzig\nTelefon 0341 904 61 12"),
    "kunde2": ("Benedikt Friese", "benedikt.friese@morgenrot-hotels.de", "Einkauf", "Morgenrot Hotels GmbH\nBrühl 38, 04109 Leipzig\nTelefon 0341 269 83 11"),
}
FILES: dict[str, list[tuple[str, str]]] = {slug: [] for slug in SLUGS}


def init_fonts():
    dirs = [Path(os.environ.get("AKTEN_FONT_DIR", "/System/Library/Fonts/Supplemental")), Path("/usr/share/fonts/truetype/msttcorefonts"), Path("/usr/share/fonts/truetype/liberation2")]
    names = {"TNR": ["Times New Roman.ttf", "times.ttf", "LiberationSerif-Regular.ttf"], "TNRB": ["Times New Roman Bold.ttf", "timesbd.ttf", "LiberationSerif-Bold.ttf"], "TNRI": ["Times New Roman Italic.ttf", "timesi.ttf", "LiberationSerif-Italic.ttf"]}
    for alias, candidates in names.items():
        path = next((d / n for d in dirs for n in candidates if (d / n).exists()), None)
        if not path:
            raise RuntimeError("Schrift fehlt; AKTEN_FONT_DIR auf Times New Roman oder Liberation Serif setzen.")
        pdfmetrics.registerFont(TTFont(alias, str(path)))
    pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNRB", italic="TNRI", boldItalic="TNRB")
    signature = Path("/System/Library/Fonts/Supplemental/Apple Chancery.ttf")
    if signature.exists():
        pdfmetrics.registerFont(TTFont("Signatur", str(signature)))
    else:
        pdfmetrics.registerFont(TTFont("Signatur", str(next(d / n for d in dirs for n in names["TNRI"] if (d / n).exists()))))


def path_for(case, name, title):
    d = ROOT / "testakten" / SLUGS[case]
    d.mkdir(parents=True, exist_ok=True)
    FILES[SLUGS[case]].append((name, title))
    return d / name


def para(text, style):
    return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)


def pdf(case, name, org, title, date, body, recipient="", ref="", table=None, widths=None, signatures=()):
    dest = path_for(case, name, title)
    normal = ParagraphStyle("Text", fontName="TNR", fontSize=11, leading=14, spaceAfter=8)
    small = ParagraphStyle("Klein", parent=normal, fontSize=9, leading=11, spaceAfter=4)
    heading = ParagraphStyle("Abschnitt", parent=normal, fontName="TNRB", spaceBefore=9, spaceAfter=12, keepWithNext=True)
    story = [para(org.split("\n")[0], ParagraphStyle("Absender", parent=normal, fontName="TNRB", fontSize=15, leading=18, spaceAfter=5)), para("\n".join(org.split("\n")[1:]), small), Spacer(1, 18)]
    if recipient:
        story += [para(recipient, normal), Spacer(1, 12)]
    story += [para(date + (" | " + ref if ref else ""), small), Spacer(1, 10), para(title, ParagraphStyle("Titel", parent=normal, fontName="TNRB", fontSize=13, leading=16, spaceAfter=16))]
    for text in body:
        if text == "PAGEBREAK":
            story.append(PageBreak())
        else:
            story.append(para(text, heading if re.match(r"^\d+(?:\.\d+)*\s", text) else normal))
    if table:
        grid = [[para(v, small) for v in row] for row in table]
        t = Table(grid, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9ecef")), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#d9d9d9")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        story += [Spacer(1, 6), t, Spacer(1, 12)]
    for sign, label in signatures:
        story.append(KeepTogether([Spacer(1, 10), para(sign, ParagraphStyle("Unterschrift", parent=normal, fontName="Signatur", fontSize=18, leading=23, spaceAfter=2)), para(label, small)]))
    def footer(c, d):
        c.saveState()
        c.setFont("TNR", 8)
        c.drawString(52, 32, org.split("\n")[0])
        c.drawRightString(A4[0] - 52, 32, f"{ref or date} | Seite {d.page}")
        c.restoreState()
    def stable_canvas(*args, **kwargs):
        kwargs["invariant"] = 1
        return canvas.Canvas(*args, **kwargs)
    doc = SimpleDocTemplate(str(dest), pagesize=A4, rightMargin=52, leftMargin=52, topMargin=45, bottomMargin=52, title=title, author=org.split("\n")[0], creator="", producer="")
    doc.build(story, onFirstPage=footer, onLaterPages=footer, canvasmaker=stable_canvas)


def docx(case, name, org, title, date, body, recipient="", signer=""):
    dest = path_for(case, name, title)
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(2.2), Cm(2.2)
    for sty in ("Normal", "Title", "Heading 1", "Heading 2"):
        d.styles[sty].font.name = "Times New Roman"
        d.styles[sty].font.size = Pt(11 if sty == "Normal" else 13)
        d.styles[sty].font.color.rgb = RGBColor(0, 0, 0)
    for fonts in d.styles.element.iter(qn('w:rFonts')):
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():
                del fonts.attrib[attr]
    for border in list(d.styles['Title'].element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    d.styles["Normal"].paragraph_format.space_after = Pt(8)
    d.styles["Normal"].paragraph_format.line_spacing = 1.08
    d.styles["Heading 1"].paragraph_format.space_after = Pt(0)
    p = d.add_paragraph()
    p.add_run(org.split("\n")[0]).bold = True
    d.add_paragraph("\n".join(org.split("\n")[1:]))
    if recipient:
        d.add_paragraph(recipient)
    d.add_paragraph(date)
    d.add_paragraph(title, "Title")
    for text in body:
        is_heading = bool(re.match(r"^\d+(?:\.\d+)*\s", text))
        p = d.add_paragraph(text, "Heading 1" if is_heading else "Normal")
        if is_heading:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
    if signer:
        d.add_paragraph(signer)
    separate_section_headings(d)
    d.core_properties.author = org.split("\n")[0]
    d.core_properties.last_modified_by = org.split("\n")[0]
    d.core_properties.title = title
    d.core_properties.created = datetime(2026, 9, 22, 12)
    d.core_properties.modified = datetime(2026, 9, 22, 12)
    d.save(dest)


def mail(case, name, sender, to, date, subject, body, cc=(), attachments=(), reply=None):
    dest = path_for(case, name, subject)
    s = CONTACTS[sender]
    msg = EmailMessage(policy=SMTP)
    msg["From"] = f"{s[0]} <{s[1]}>"
    msg["To"] = ", ".join(f"{CONTACTS[k][0]} <{CONTACTS[k][1]}>" for k in to)
    if cc:
        msg["Cc"] = ", ".join(f"{CONTACTS[k][0]} <{CONTACTS[k][1]}>" for k in cc)
    msg["Date"] = datetime.fromisoformat(date)
    msg["Subject"] = subject
    msg["Message-ID"] = f"<{dest.stem}@{s[1].split('@')[1]}>"
    if reply:
        msg["In-Reply-To"] = reply
        msg["References"] = reply
    if attachments:
        msg["X-Attachments"] = ", ".join(attachments)
    msg.set_content(body.strip() + "\n\n" + s[0] + ("\n" + s[2] if s[2] else "") + "\n" + s[3] + "\n", charset="utf-8")
    dest.write_bytes(msg.as_bytes())


def csv_file(case, name, title, rows):
    dest = path_for(case, name, title)
    with dest.open("w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f, delimiter=";").writerows(rows)


def txt(case, name, title, text):
    path_for(case, name, title).write_text(text.strip() + "\n", encoding="utf-8")


def screen_fonts(size):
    from akten_build_runtime import screen_font
    return screen_font(size)


def machine_screen():
    from PIL import Image, ImageDraw
    dest = path_for(0, '26_px12_betriebsanzeige_20-08.png', 'Betriebsanzeige PX12 am 20. August 2026')
    im = Image.new('RGB', (1440, 1000), '#edf0ed')
    d = ImageDraw.Draw(im)
    def text(x, y, value, size=26, color='#17241e'):
        d.text((x, y), value, font=screen_fonts(size), fill=color)
    d.rectangle((0, 0, 1440, 108), fill='#263d33')
    text(38, 24, 'PX12 Betriebsanzeige', 38, 'white')
    text(985, 34, '20.08.2026  11:00:08', 27, 'white')
    text(38, 134, 'Maschine PX12-260071', 28)
    text(890, 134, 'Automatik   /   Bereit', 28)
    d.line((38, 184, 1400, 184), fill='#9daaa2', width=2)
    text(38, 213, 'Rezept 04   Kaffee 450 g', 32)
    text(38, 268, 'Beutel: SF-PET12-PE80     140 × 80 × 260 mm', 27)
    text(38, 316, 'Materialcharge: F260728B', 27)
    blocks = [(38,385,'Taktvorgabe','5,5','Beutel/min'), (509,385,'Temperatur Soll','175','Grad Celsius'), (980,385,'Siegelzeit','1,0','Sekunden')]
    for x,y,label,value,unit in blocks:
        d.rectangle((x,y,x+420,y+220),fill='white',outline='#aeb9b1',width=2)
        text(x+22,y+22,label,26)
        text(x+22,y+72,value,66)
        text(x+22,y+166,unit,25)
    text(38, 657, 'Zykluszähler seit Rücksetzen: 330', 32)
    text(38, 713, 'Letztes Rücksetzen: 20.08.2026 09:58:11', 26)
    text(38, 771, 'Haubenkontakt geschlossen     Not-Halt frei', 26)
    d.rectangle((0, 846, 1440, 1000), fill='#dbe2dc')
    text(38, 873, 'Meldungen: Keine aktive Störung', 29)
    text(38, 928, 'Nahtprüfung extern. Prüfergebnisse werden nicht in der Steuerung gespeichert.', 24)
    im.save(dest)


def event_screen(case, name, title, rows):
    from PIL import Image, ImageDraw
    dest = path_for(case, name, title)
    im = Image.new('RGB', (1860, 1230), '#f4f5f7')
    d = ImageDraw.Draw(im)
    def text(x, y, value, size=23, color='#1e2429'):
        d.text((x,y),value,font=screen_fonts(size),fill=color)
    d.rectangle((0,0,1860,92),fill='#343c44')
    text(30,25,'Mühlentor Sensortechnik | Geräteverwaltung',34,'white')
    text(30,125,'MT-L019   Ereignisverlauf',32)
    text(1230,130,'Oskar Seitz | 17.09.2026 11:10',25)
    text(30,183,'Filter: 10.09.2026 bis 17.09.2026    Benutzer t.hesse / o.seitz',25)
    d.rectangle((30,245,1830,299),fill='#dce2e7')
    xs = [42,315,525,820,1365]
    for x,label in zip(xs,['Zeit CEST','System','Ereignis','Objekt','Ergebnis']):
        text(x,259,label,24)
    display = [
        ['10.09. 15:59:12','CRM','Anmeldung','t.hesse / CRM-9951','Erfolgreich'],
        ['10.09. 16:02:38','CRM','Export','Kunden mit Umsatz 2025','412 Datensätze'],
        ['10.09. 16:02:42','MT-L019','Datei erstellt','kunden_20260910.csv','188.416 Byte'],
        ['10.09. 16:04:05','MT-L019','USB verbunden','32 GB / Kennung 82A119','Angeschlossen'],
        ['10.09. 16:04:17','MT-L019','Datei geöffnet','kunden_20260910.csv','Tabellenanwendung'],
        ['10.09. 16:08:44','MT-L019','USB getrennt','Kennung 82A119','Ordnungsgemäß'],
        ['11.09. 09:17:28','Teamlaufwerk','Datei gespeichert','Budget/serienkunden_2027.xlsx','28.416 Byte'],
        ['11.09. 09:21:07','Teamlaufwerk','Datei geöffnet','n.berg / serienkunden_2027.xlsx','Lesezugriff'],
        ['17.09. 10:15:21','MT-L019','Bestand aufgenommen','kunden_20260910.csv','188.416 Byte'],
        ['17.09. 10:34:02','Verzeichnisdienst','Zugang gesperrt','t.hesse','Web und Fernzugriff'],
        ['17.09. 11:08:00','IT-Ausgabe','Exportvermerk','USB-Dateikopiervorgänge','Nicht protokolliert'],
    ]
    assert len(display) == len(rows)-1
    for i,row in enumerate(display):
        y = 300+i*65
        if i%2 == 0:
            d.rectangle((30,y,1830,y+65),fill='white')
        for x,value in zip(xs,row):
            text(x,y+19,value,22)
    text(30,1056,'11 Ereignisse | Anzeige einschließlich IT-Ausgabevermerk EV-411',25)
    text(30,1112,'Zielpfade auf privaten USB-Datenträgern werden nicht protokolliert.',25)
    text(30,1164,'Datei auf MT-L019: Downloads/kunden_20260910.csv',25)
    im.save(dest)


def leipzig():
    mail(0, "01_ott_an_stahl_22-09-2026.eml", "ott", ["stahl"], "2026-09-22T08:12:00+02:00", "PX12 / Nordwerk verlangt 28.560 EUR und kündigt Einzug an", """Sehr geehrte Frau Dr. Stahl,

bitte übernehmen Sie die Korrespondenz mit Nordwerk. Wir brauchen die Anlage für unsere 450-g-Beutel, nicht für die kleinen Probiertüten. Bis heute erreichen wir die zugesagten acht Beutel in der Minute nicht. Die Rechnung und der Brief vom 8. September liegen bei. Für Freitag ist darin ein Einzug angekündigt. Leonie findet kein Mandat für Nordwerk.

Wir haben die Anzahlung geleistet und wollen nicht einfach eine zweite Maschine bestellen. Zwei Kunden warten auf zusammen 4.200 Beutel. Ich möchte vor einer Antwort wissen, welche Angaben Ihnen noch fehlen und ob Sie mit Herrn Behn einen verbindlichen Reparaturtermin erreichen können. Bitte noch keine Ersatzbestellung auslösen. Das Angebot aus Chemnitz habe ich nur angefragt.

Maja hat den Betrieb am 18. August begleitet. Ich war erst gegen Mittag da und habe nichts als endgültige Abnahme unterschrieben. Die Original-Mail mit dem Angebot vom 4. März fehlt in unserem alten Postfach; im Auftragsordner liegt die beigefügte AGB-Fassung. Ob damals weitere Anlagen dabei waren, weiß ich nicht. Herrn Haas habe ich am 31. August nicht mehr erreicht.

Freundliche Grüße""", attachments=("17_rechnung_nw-261184.pdf", "18_restforderung_08-09-2026.pdf"))
    pdf(0, "02_auftrag_px12_12-03-2026.pdf", NORD, "Auftrag PX12 mit Einrichtung am Einsatzort", "Leipzig / Hamburg, 12.03.2026", [
        "Zwischen der Nordwerk Verpackungstechnik GmbH und der Elsterbogen Rösterei GmbH, Weißenfelser Straße 42, 04229 Leipzig, wird der Auftrag NW-260312-47 vereinbart.",
        "1 Lieferumfang und Einrichtung",
        "Nordwerk liefert eine Beutelverschließmaschine PX12 mit Zuführband, höhenverstellbarer Beutelführung und Siegelbackensatz 180 mm. Nordwerk richtet die Führung für stehende, von Hand vorbefüllte Ventilbeutel mit 450 g Röstkaffee ein. Füllwaage und Dosierung sind nicht Bestandteil des Auftrags. Eine Einweisung von zwei Beschäftigten und ein gemeinsamer Lauf mit dem Kundenmaterial am Aufstellort sind im Preis enthalten.",
        "2 Vereinbarte Leistung",
        "Mit den Beuteln Saaleflex SF-PET12-PE80, Format 140 × 80 × 260 mm, beträgt die vereinbarte Leistung acht verschlossene Beutel je Minute. Nach dem Aufwärmen wird ein ununterbrochener Lauf von 30 Minuten mit 450-g-Füllung durchgeführt. Die Nähte müssen bei der vereinbarten Druckprobe geschlossen bleiben. Bedienung, Takt und Nahtkontrolle werden im Inbetriebnahmeblatt dokumentiert. Die endgültige Abnahme wird nach diesem Lauf gesondert bestätigt. Das Produktblatt PX12, Stand 02.03.2026, ist beigefügt.",
        "3 Termin und Vergütung",
        "Die Lieferung erfolgt am 17.08.2026 frei Verwendungsstelle Leipzig. Die Einrichtung ist für den Folgetag vorgesehen. Der Festpreis einschließlich Einrichtung und Einweisung beträgt 36.000,00 EUR netto zuzüglich 19 Prozent Umsatzsteuer von 6.840,00 EUR, insgesamt 42.840,00 EUR. Ein Drittel, 14.280,00 EUR brutto, wird nach Auftragserteilung berechnet. Der Restbetrag von 28.560,00 EUR brutto wird nach Lieferung und Inbetriebnahme berechnet und ist binnen 14 Tagen nach Rechnungsdatum zahlbar.",
        "4 Weitere Vereinbarungen",
        "Stromanschluss 230 V und trockene, ebene Stellfläche stellt Elsterbogen bereit. Nordwerk liefert mit Schutzhauben, deutschsprachiger Betriebsanleitung und Dokumentation. Reparaturtermine werden mit der Produktionsleitung abgestimmt. Änderungen der Produktauslegung werden vor Ausführung schriftlich vereinbart. Als Gerichtsstand ist Leipzig vereinbart. Ergänzend gelten die mit dem Angebot NW-260304-47 übersandten Lieferbedingungen, Stand 01.02.2026. Individuelle Vereinbarungen dieses Auftrags gehen vor. Jede Partei erhält eine unterschriebene Ausfertigung."
    ], ref="NW-260312-47", signatures=(("Friederike Ott", "Elsterbogen Rösterei GmbH | Leipzig, 12.03.2026"), ("Hendrik Behn", "Nordwerk Verpackungstechnik GmbH | Hamburg, 12.03.2026")))
    pdf(0, "03_lieferbedingungen_01-02-2026.pdf", NORD, "Allgemeine Lieferbedingungen", "Stand 01.02.2026", [
        "1 Geltung und Auftrag",
        "Diese Bedingungen gelten für Lieferungen und Serviceleistungen an Unternehmer. Abweichende Bedingungen des Bestellers werden nur Vertragsbestandteil, wenn Nordwerk sie schriftlich bestätigt. Angebote gelten 30 Tage. Umfang und Ausführung ergeben sich aus der Auftragsbestätigung und den ausdrücklich einbezogenen technischen Unterlagen.",
        "2 Lieferung und Mitwirkung",
        "Der Besteller stellt geeignete Anschlüsse und das abgestimmte Verpackungsmaterial rechtzeitig bereit. Verzögerungen aufgrund fehlender Mitwirkung werden mitgeteilt. Teillieferungen sind nur zulässig, soweit sie für den Besteller selbständig nutzbar sind. Bei erkennbaren Transportschäden soll der Besteller den Frachtführer hinzuziehen und Nordwerk unverzüglich informieren.",
        "3 Zahlung und Aufrechnung",
        "Rechnungen sind ohne Abzug binnen 14 Tagen zahlbar. Eine Aufrechnung des Bestellers ist ausgeschlossen, sofern die Gegenforderung nicht unbestritten oder rechtskräftig festgestellt ist. Der Besteller darf Zahlungen wegen anderer Geschäfte nicht zurückhalten. Eine Einzugsermächtigung bedarf einer gesonderten Erklärung. Bis zur vollständigen Bezahlung bleibt die gelieferte Maschine Eigentum von Nordwerk.",
        "4 Beanstandungen und Service",
        "Beanstandungen sollen Maschinennummer, eingesetztes Material, Einstellungen und beobachtetes Fehlerbild enthalten. Nordwerk erhält Gelegenheit zur Untersuchung und zur Beseitigung festgestellter Fehler. Verschleißteile sind ausgeschlossen, soweit der Ausfall auf gewöhnlicher Abnutzung beruht. Kosten für Einsätze wegen ungeeigneter Verbrauchsmaterialien werden nach vorheriger Mitteilung berechnet. Für Schäden aus vorsätzlichem oder grob fahrlässigem Verhalten sowie für Schäden an Leben, Körper oder Gesundheit gelten die gesetzlichen Bestimmungen.",
        "5 Erfüllungsort und Gerichtsstand",
        "Erfüllungsort für Zahlungen ist Hamburg. Für alle Streitigkeiten mit Kaufleuten wird Hamburg als ausschließlicher Gerichtsstand vereinbart. Es gilt deutsches Recht. Die Unwirksamkeit einer einzelnen Bestimmung berührt die übrigen Vereinbarungen nicht."
    ], ref="LB 02/2026")
    pdf(0, "04_datenblatt_px12_450g.pdf", NORD, "PX12 Technische Daten und Rüstwerte", "02.03.2026", [
        "Ausführung für Auftrag NW-260304-47. Die Maschine verschließt vorbefüllte Standbodenbeutel. Die Füllmenge von 450 g bezieht sich auf geröstete ganze Kaffeebohnen. Freier Kopfraum unterhalb der Naht: mindestens 35 mm. Das Entgasungsventil muss außerhalb der Siegelzone liegen.",
        "Bei Material SF-PET12-PE80 liegt das Einstellfenster bei 160 bis 180 Grad Celsius und 0,8 bis 1,2 Sekunden Siegelzeit. Die tatsächliche Backentemperatur ist nach 15 Minuten Aufwärmzeit mit einem Kontaktfühler zu kontrollieren. Die Displayanzeige allein ersetzt diese Kontrolle nicht.",
        "Der PX12-Taktwert bezeichnet vollständige Beutelzyklen pro Minute. Für die Auftragsausführung sind acht Beutel pro Minute bei 450-g-Befüllung vereinbart. Höhere Katalogwerte für ungefüllte Musterbeutel gelten hierfür nicht. Staub in der Nahtzone ist vor dem Einlauf zu entfernen. Papierverbunde benötigen ein gesondertes Backenprofil."
    ], ref="PX12 / Ausgabe 03-2026", table=[['Merkmal', 'Ausführung'], ['Netz / Leistung', '230 V / 50 Hz / 1,8 kW'], ['Abmessungen / Masse', '1.240 × 670 × 1.180 mm / 146 kg'], ['Beutelformat', '140 × 80 × 260 mm; Nahtbreite 10 mm'], ['Material', 'PET 12 µm / PE 80 µm; Artikel SF-PET12-PE80'], ['Druckprobe', 'Jeder Beutel 5 Sekunden mit 2-kg-Prüfplatte belasten; offene Naht zählt als Fehler'], ['Sicherungen', 'Not-Halt, Haubenkontakt, temperaturabhängige Freigabe']], widths=[140, 350])
    pdf(0, "05_anzahlungsrechnung_nw-260341.pdf", NORD, "Anzahlungsrechnung NW-260341", "16.03.2026", [
        "Sehr geehrte Frau Ott,", "für den Auftrag NW-260312-47 berechnen wir die vereinbarte Anzahlung. Der voraussichtliche Leistungszeitpunkt ist August 2026. Bitte überweisen Sie bis 23.03.2026 unter Angabe der Rechnungsnummer auf das unten genannte Konto.", "Umsatzsteuer-ID DE328741652. Bankverbindung: Hanse Geschäftsbank, IBAN DE20 2005 0000 1234 5678 90. Zahlungen werden auf die Schlussrechnung angerechnet.", "Mit freundlichen Grüßen\nLea Kuhl\nBuchhaltung"
    ], recipient=ELSTER.split('\n')[0] + '\nFrau Friederike Ott\nWeißenfelser Straße 42\n04229 Leipzig', ref="Kunde 1048", table=[['Position', 'EUR'], ['Anzahlung netto', '12.000,00'], ['Umsatzsteuer 19 Prozent', '2.280,00'], ['Zahlbetrag brutto', '14.280,00']], widths=[360, 130])
    pdf(0, "06_kontoauszug_anzahlung.pdf", "Elsterbank Leipzig\nFirmenkundenservice, Postfach 10 14 72, 04014 Leipzig\nTelefon 0341 993 20 0", "Umsatzanzeige Geschäftskonto", "Abruf 21.09.2026, 10:04 Uhr", [
        "Kontoinhaberin: Elsterbogen Rösterei GmbH. Geschäftskonto IBAN DE42 8605 5592 1100 4433 21. Auswahl: Buchung vom 20.03.2026. Alle Beträge in EUR.", "Die Überweisung wurde am 19.03.2026 um 16:22 Uhr von Friederike Ott freigegeben und am nächsten Geschäftstag ausgeführt. Die Umsatzanzeige enthält ausschließlich die ausgewählte Buchung."
    ], ref="Umsatz 2026-03-20/184", table=[['Buchung / Wertstellung', 'Empfänger und Verwendungszweck', 'Belastung'], ['20.03.2026 / 20.03.2026', 'Nordwerk Verpackungstechnik GmbH\nDE20 2005 0000 1234 5678 90\nNW-260341 / NW-260312-47', '14.280,00']], widths=[110, 280, 100])
    pdf(0, "07_lieferschein_17-08-2026.pdf", NORD, "Lieferschein LS-260817-93", "17.08.2026, 09:35 Uhr", [
        "Empfangsort: Elsterbogen Rösterei GmbH, Weißenfelser Straße 42, Halle 2, 04229 Leipzig. Angeliefert durch Frachtfahrer Erik Linde. Auftrag NW-260312-47, Maschinennummer PX12-260071.",
        "Übergeben wurden ein PX12 mit Zuführband und Schutzhaube, ein Beutelführungssatz 450 g, Betriebsanleitung Ausgabe 03/2026 und eine Werkzeugtasche. Stromanschluss wird am 18. August gemeinsam hergestellt. Die Palette ist äußerlich unbeschädigt.",
        "Empfangsvermerk Maja Brandt: Vollständigkeit der oben aufgeführten Packstücke bestätigt. Maschine heute nicht in Betrieb genommen. Keine Funktionsfreigabe. Bitte morgen den kleineren Führungswinkel mitbringen; dieser liegt nicht in der Werkzeugtasche."
    ], ref="LS-260817-93", signatures=(("Maja Brandt", "Empfang für Elsterbogen Rösterei GmbH, 17.08.2026"), ("Erik Linde", "Fahrer, 17.08.2026")))
    raw = [['Messung','Datum','Beginn','Minuten','Beutel_450g','Naht_offen','Display_Grad_C','Siegelzeit_s','Takt_Anzeige','Charge','Bediener']]
    leaks = [4,5,4,3,4,3,5,4,5,2,3,2,3,4,3,4,3,4]
    for i in range(18):
        day = [18,19,20][i // 6]
        raw.append([f'P{i+1:02}', f'2026-08-{day}', f'{10+i%6//6:02}:{(i%6)*10:02}',10,55,leaks[i],170 if i<12 else 175,'1,0', '5,5','F260728B','M. Brandt' if i<12 else 'J. Voigt'])
    csv_file(0, "08_produktionsbuch_px12.csv", "Produktionsbuch PX12 vom 18. bis 20. August", raw)
    pdf(0, "09_inbetriebnahme_18-08-2026.pdf", NORD, "Inbetriebnahmeblatt PX12-260071", "18.08.2026, 08:30 bis 11:20 Uhr", [
        "Anwesend waren Monteur Florian Lüders und Maja Brandt, Produktionsleitung Elsterbogen. Der fehlende Führungswinkel wurde montiert. Netz, Haubenschalter und Not-Halt wurden betätigt; die Antriebe stoppten. Frau Brandt wurde in Reinigung und Temperaturwahl eingewiesen.",
        "Der Lauf von 10:00 bis 10:30 Uhr ergab 165 Beutel mit je 450 g. Die Taktanzeige stand auf 5,5. Bei Erhöhung auf 8,0 kippten zwei Beutel am Übergang zur Führung. Nach Rückstellung auf 5,5 lief die Zuführung ohne Kippen. Bei 13 Beuteln öffnete sich die Naht unter der Prüfplatte, überwiegend am rechten Rand. Diese Beutel wurden ausgesondert. Produktionsbuch: P01 bis P03.",
        "Material laut Karton: Saaleflex SF-PET12-PE80, Charge F260728B. Display 170 Grad Celsius, Siegelzeit 1,0 Sekunden. Die Backentemperatur wurde nicht separat gemessen; der Kontaktfühler lag im anderen Servicefahrzeug. Florian Lüders kündigte einen weiteren Termin zur Einstellung der Führung und des Backendrucks an.",
        "Vermerk Maja Brandt: Lauf und Einweisung bestätigt. Endgültige Abnahme bleibt offen; Leistung und Nähte müssen nochmals gemeinsam vorgeführt werden. Die Maschine darf für begrenzte betriebliche Läufe unter vollständiger Nahtkontrolle genutzt werden."
    ], ref="IB-260818-71", signatures=(("Florian Lüders", "Service Nordwerk"), ("Maja Brandt", "Produktion Elsterbogen; vorstehender Vermerk bei Unterzeichnung ergänzt")))
    mail(0, "10_email_probelauf_abnahme_offen.eml", "brandt", ["haas"], "2026-08-18T13:46:00+02:00", "PX12 läuft im Probelauf, endgültige Abnahme offen", """Guten Tag Herr Haas,

die PX12 läuft im Probelauf, die endgültige Abnahme ist offen. Das haben Herr Lüders und ich auch auf seinem Blatt festgehalten. Die 165 Beutel im halbstündigen Lauf ergeben 5,5 pro Minute, nicht acht. Bei 13 Nähten konnten wir die rechte Ecke unter der Platte öffnen. Die betroffenen Beutel sind getrennt in Kiste 3.

Wir verwenden die 450-g-Beutel aus dem besprochenen PET/PE-Verbund. Die wenigen durchgelaufenen Beutel brauchen wir heute, aber jede Naht wird einzeln geprüft. Bitte schicken Sie Herrn Lüders mit Kontaktfühler und dem anderen Führungsanschlag. Frau Ott bittet um zwei Terminvorschläge noch in dieser Woche.

Freundliche Grüße""", cc=("ott",))
    docx(0, "11_maengelanzeige_20-08-2026.docx", ELSTER, "Beanstandung der PX12 und Bitte um Beseitigung", "Leipzig, 20.08.2026 | Auftrag NW-260312-47", [
        "Sehr geehrter Herr Behn,",
        "die am 17. August gelieferte PX12 erreicht bei unseren vereinbarten 450-g-Beuteln bisher nur 5,5 statt acht Beutel je Minute. Im gemeinsamen Lauf am 18. August waren 13 von 165 Nähten undicht. Auch die anschließenden Läufe zeigen offene Nähte. Die Ware können wir so nicht versenden. Bei höherem Takt kippen Beutel am Führungsübergang.",
        "Bitte beseitigen Sie diese Probleme bis einschließlich 31.08.2026 und führen Sie anschließend den vereinbarten halbstündigen Lauf mit uns durch. Den Betrieb können wir am 26. und 27. August jeweils ab 8 Uhr für Sie freihalten. Bitte bestätigen Sie den Termin bis morgen. Ungeprüfte Beutel werden nicht an Kunden herausgegeben.",
        "Wir halten am Auftrag fest und stellen die Maschine für den Service bereit. Eine endgültige Abnahme haben wir nicht erklärt. Die Rechnung vom 19. August haben wir erhalten; den Restbetrag geben wir wegen der genannten Punkte derzeit nicht frei. Die vereinbarte Anzahlung wurde am 20. März überwiesen.",
        "Mit freundlichen Grüßen"
    ], recipient="Nordwerk Verpackungstechnik GmbH\nHerrn Hendrik Behn\nBillbrookdeich 168\n22113 Hamburg", signer="Friederike Ott\nGeschäftsführerin")
    mail(0, "12_nordwerk_behauptung_folie.eml", "haas", ["ott"], "2026-08-21T09:18:00+02:00", "Ihre Beanstandung vom 20. August / Verpackungsmaterial", """Sehr geehrte Frau Ott,

nach Rücksprache mit Herrn Lüders vermuten wir, dass Ihre Beutel einen Papieranteil haben. Für solche Verbunde gilt das PET/PE-Rüstfenster nicht. In diesem Fall erklären sich die offenen Ecken aus dem Material und nicht aus der Maschine. Im Lagerfoto, das Herr Lüders mir auf dem Telefon gezeigt hat, sehe ich braune Kartons und matte Beutel; eine Artikelnummer kann ich darauf nicht erkennen.

Wir lassen den Einsatz am 26. August disponieren. Halten Sie bitte ungefüllte Beutel und die Kartonetiketten bereit. Herr Lüders soll zuerst die Temperatur prüfen. Wenn ungeeignete Folie vorliegt, möchten wir die Fahrt berechnen. Eine Kostenfreigabe von Ihnen liegt bislang nicht vor. Die Rechnung über die Maschine bleibt bei uns offen.

Mit freundlichen Grüßen""", cc=("brandt",))
    mail(0, "13_materialempfehlung_saaleflex.eml", "folie", ["brandt", "haas"], "2026-03-09T11:06:00+01:00", "PX12 / Empfehlung SF-PET12-PE80 für 450 g Kaffee", """Guten Tag Frau Brandt, guten Tag Herr Haas,

für die abgestimmte PX12-Ausführung empfehlen wir den Artikel SF-PET12-PE80 in 140 × 80 × 260 mm. Der Verbund besteht aus 12 µm PET und 80 µm PE, ohne Papierlage. Die Siegelzone oben bleibt unbedruckt. Das Ventil sitzt 65 mm unter der Oberkante; bei 450 g ganzen Bohnen bleibt ausreichend freier Kopfraum.

Unser Materialfenster beträgt 160 bis 180 Grad Celsius bei 0,8 bis 1,2 Sekunden Kontaktzeit. Bitte beginnen Sie bei 170 Grad und 1,0 Sekunden. Maßgeblich ist die Temperatur an der Backenoberfläche. Einseitig offene Nähte sollten Sie auch auf ungleichmäßigen Druck und Falten im Übergang prüfen. Eine Freigabe Ihrer konkreten Maschine können wir aus der Ferne nicht erteilen.

Die erste Serienlieferung ist für August vorgesehen. Die genaue Fertigungscharge teilen wir auf dem Lieferschein mit. Muster aus dem Februar tragen noch die Kennung M260219 und dürfen nicht mit der Seriencharge verwechselt werden.

Freundliche Grüße""")
    pdf(0, "14_saaleflex_lieferschein_charge.pdf", "Saaleflex Folien GmbH\nDieselstraße 19, 06112 Halle (Saale)\nTelefon 0345 681 27 0 | versand@saaleflex.de", "Lieferschein SF-260812-208", "12.08.2026", [
        "Lieferung an Elsterbogen Rösterei GmbH, Weißenfelser Straße 42, 04229 Leipzig. Ihre Bestellung EB-260706-12 vom 06.07.2026. Auslieferung mit Stückgut am 12.08.2026; Wareneingang am 13.08.2026 um 10:15 Uhr.",
        "Alle 20 Kartons tragen die Charge F260728B. Artikel SF-PET12-PE80, PET 12 µm / PE 80 µm, ohne Papier. Standbodenbeutel mit Ventil, 140 × 80 × 260 mm, für 450 g Kaffeebohnen. Je Karton 500 Beutel; Gesamtmenge 10.000 Stück.",
        "Wareneingang Maja Brandt: 20 Kartons gezählt, trocken und geschlossen. Kartons 01 und 02 am 18.08. an der PX12 geöffnet. Etikett von Karton 01 in Produktionsmappe übernommen: SF-PET12-PE80 / F260728B / 01-20 / 500 St."
    ], ref="SF-260812-208", signatures=(("Maja Brandt", "Wareneingang am 13.08.2026; Entnahmevermerk am 18.08.2026"),))
    pdf(0, "15_servicebericht_26-08-2026.pdf", NORD, "Servicebericht PX12", "26.08.2026, 08:10 bis 10:05 Uhr", [
        "Maschine PX12-260071, Auftrag NW-260312-47. Monteur Florian Lüders. Ansprechpartnerin vor Ort Maja Brandt. Einsatzgrund: Nahtöffnung rechts und Kippen bei Takt 8,0.",
        "Führung um 3 mm versetzt. Rechte Siegelbacke gereinigt; an der Innenseite dunkler Abrieb. Display 175 Grad Celsius. Kontaktmessung links 174 Grad, rechts 158 Grad. Kontaktfühler KF-19; letzter Prüfaufkleber 05/2025. Messung nach Öffnen der Haube; Zeitabstand zwischen den beiden Messpunkten nicht notiert. Rechte Heizpatrone elektrisch durchgängig. Eine Vergleichsmessung bei geschlossener Haube war mit diesem Fühler nicht möglich.",
        "Backendruck per Stellschraube erhöht, ohne Messuhr. Lauf 09:25 bis 09:35 Uhr: 60 Beutel, davon drei rechts offen. Bei Takt 8,0 nochmals zwei Kipper; für den Lauf Takt auf 6,0 gestellt. Eingesetzte Beutel laut Frau Brandt unverändert. Kartonetikett nicht mitgenommen. Keine Proben zur Werkstatt zurückgeführt.",
        "Ersatz des rechten Heizpatronensatzes und ein weiterer Versuch mit Druckmessung werden vorgesehen. Lagerbestand wird in Hamburg geklärt. Frau Brandt verlangt weiterhin acht Beutel je Minute ohne Nahtöffnung. Sie bestätigt die Anwesenheitszeit, nicht den Abschluss der Arbeiten."
    ], ref="SV-260826-114", signatures=(("Florian Lüders", "Servicetechniker"), ("Maja Brandt", "Anwesenheit bestätigt; Beanstandung besteht fort")))
    mail(0, "16_terminbestaetigung_26-08.eml", "haas", ["brandt"], "2026-08-24T14:25:00+02:00", "Bestätigt: Service am 26. August ab 8 Uhr", """Guten Tag Frau Brandt,

Herr Lüders kommt am Mittwoch, 26. August, zwischen 8 und 8:30 Uhr. Er bringt einen Kontaktfühler und einen verstellbaren Führungsanschlag mit. Der Wagen muss bis 11 Uhr weiter. Bitte halten Sie mindestens 300 Beutel aus der derzeit verwendeten Charge und 140 kg Kaffee bereit; bereits ausgesonderter Kaffee kann für die Einstellversuche genutzt werden.

Wir arbeiten zunächst an Führung und Naht. Sollte ein Teil fehlen, meldet sich Herr Lüders vor Ort bei mir. Eine gesonderte Kostenvereinbarung treffen wir mit dieser Terminbestätigung nicht. Ihre im Brief genannte Frist habe ich an die Geschäftsleitung weitergegeben.

Freundliche Grüße""", cc=("ott",))
    pdf(0, "17_rechnung_nw-261184.pdf", NORD, "Rechnung NW-261184", "19.08.2026", [
        "Sehr geehrte Frau Ott,", "wir berechnen die PX12 einschließlich Beutelführung, Einrichtung und Einweisung gemäß Auftrag NW-260312-47. Lieferdatum 17.08.2026, Inbetriebnahmebesuch 18.08.2026. Die Anzahlung ist am 20.03.2026 eingegangen und wird unten abgesetzt.",
        "Bitte zahlen Sie den verbleibenden Betrag bis 02.09.2026 auf IBAN DE20 2005 0000 1234 5678 90, Hanse Geschäftsbank. Umsatzsteuer-ID DE328741652. Verwendungszweck NW-261184 / Kunde 1048.", "Mit freundlichen Grüßen\nLea Kuhl\nBuchhaltung"
    ], recipient="Elsterbogen Rösterei GmbH\nWeißenfelser Straße 42\n04229 Leipzig", ref="NW-261184", table=[['Position', 'Netto EUR', 'USt EUR', 'Brutto EUR'], ['PX12 einschließlich Einrichtung', '36.000,00', '6.840,00', '42.840,00'], ['Abzüglich Anzahlung NW-260341', '-12.000,00', '-2.280,00', '-14.280,00'], ['Verbleibender Betrag / 19 Prozent', '24.000,00', '4.560,00', '28.560,00']], widths=[220,90,90,90])
    pdf(0, "18_restforderung_08-09-2026.pdf", NORD, "Offener Betrag aus Rechnung NW-261184", "Hamburg, 08.09.2026", [
        "Sehr geehrte Frau Ott,", "auf unserem Kundenkonto 1048 stehen weiterhin 28.560,00 EUR offen. Die Rechnung war am 2. September zahlbar. Die Anlage wurde geliefert, eingerichtet und wird von Ihnen genutzt. Wir erwarten den vollständigen Ausgleich.",
        "Unsere Technik hat Ihnen die Eignung der Folie als Ursache genannt. Der weitere Besuch dient aus unserer Sicht der Unterstützung Ihres Betriebs. Die Zahlungsbedingungen lassen eine Aufrechnung nur mit unbestrittenen oder rechtskräftig festgestellten Gegenforderungen zu. Bitte überweisen Sie bis 18.09.2026. Anderenfalls sehen wir den Einzug vom hinterlegten Konto mit Endziffern 3321 am 25.09.2026 vor. Als Einzugsreferenz ist NW1048 eingetragen.",
        "Bitte melden Sie sich unmittelbar bei mir, wenn Sie bereits überwiesen haben.", "Mit freundlichen Grüßen"
    ], recipient="Elsterbogen Rösterei GmbH\nFrau Friederike Ott\nWeißenfelser Straße 42\n04229 Leipzig", ref="Kunde 1048 / NW-261184", signatures=(("Hendrik Behn", "Geschäftsführer"),))
    mail(0, "19_buchhaltung_lastschriftkonto.eml", "pohl", ["ott"], "2026-09-09T10:32:00+02:00", "Nordwerk / Konto 3321 und Mandatsordner", """Hallo Friederike,

die Endziffern 3321 gehören zu unserem Geschäftskonto DE42 8605 5592 1100 4433 21. Nordwerk kennt es von der Überweisung über 14.280 EUR. Im Mandatsordner liegen Strom, Miete und Saaleflex, aber kein Nordwerk-Mandat. Im Kundenblatt von Nordwerk steht NW1048; auf unserer Auftragsausfertigung ist kein Einzugsfeld.

Ich habe heute um 10:05 Uhr mit Frau Kunze beim Firmenkundenservice gesprochen. Sie sieht bis jetzt keinen vorgemerkten Nordwerk-Einzug. Die Bank braucht für eine gezielte Sperre auch die Gläubigerkennung. Die steht im Schreiben nicht. Ich habe noch keine Sperre beantragt und keine Rücklastschrift veranlasst.

Bis einschließlich des heutigen Abrufs ist außer der Anzahlung nichts an Nordwerk abgegangen. Die 28.560 EUR sind in unserer Zahlungsdatei vorgemerkt, aber von dir nicht freigegeben. Soll ich Frau Kuhl um eine Kopie des angeblichen Mandats bitten?

Viele Grüße""")
    mail(0, "20_nordwerk_ersatzteil_termin.eml", "haas", ["ott", "brandt"], "2026-08-31T16:48:00+02:00", "PX12 / Heizpatrone und nächster Besuch", """Guten Tag Frau Ott,

den für heute besprochenen Rückruf habe ich nicht geschafft. Der rechte Heizpatronensatz ist nicht im Hamburger Lager. Der Zulieferer nennt derzeit den 21. September. Ich kann Ihnen einen Besuch am 28. September um 8:30 Uhr anbieten, sobald das Teil eingetroffen ist. Eine feste Zusage zur Teilelieferung habe ich noch nicht.

Die Materialfrage ist aus unserer Sicht weiterhin offen. Bitte bewahren Sie die angebrochenen Kartons auf. Einen gegengezeichneten Abschlussbericht haben wir nicht. Herrn Behn habe ich die Messwerte des letzten Besuchs weitergegeben. Ob er die Rechnung bis zum nächsten Versuch ruhend stellt, kann ich Ihnen nicht zusagen.

Freundliche Grüße""")
    pdf(0, "21_angebot_ersatzmaschine_16-09.pdf", "Muldenpack Maschinenbau GmbH\nAltchemnitzer Straße 74, 09120 Chemnitz\nTelefon 0371 419 82 0 | vertrieb@muldenpack.de", "Angebot MP-260916-33", "16.09.2026", [
        "Sehr geehrte Frau Ott,", "auf Ihre telefonische Anfrage bieten wir eine MP9-V mit Führung für 450-g-Standbodenbeutel an. Für den Artikel SF-PET12-PE80 werden neun Beutel je Minute angeboten. Vor verbindlicher Bestellung bitten wir um 200 leere Beutel und 20 kg Kaffee für den Lauf in unserem Werk.",
        "Der Preis einschließlich Anlieferung, Einrichtung und Einweisung in Leipzig beträgt 39.800,00 EUR netto, 7.562,00 EUR Umsatzsteuer und 47.362,00 EUR brutto. Zahlung: 30 Prozent nach Bestellung, Rest nach Vorführung am Einsatzort. Ein Gerät kann bis 24.09. unverbindlich reserviert werden. Nach Bestellung bis zu diesem Datum ist Lieferung voraussichtlich ab 05.10.2026 möglich.",
        "Das Angebot gilt bis 24.09.2026. Die Inzahlungnahme einer PX12 ist nicht enthalten. Eine Miete oder Überbrückungsmaschine können wir im September nicht anbieten. Eine Bestellung ist bei uns noch nicht eingegangen.", "Mit freundlichen Grüßen"
    ], recipient="Elsterbogen Rösterei GmbH\nFrau Friederike Ott\nWeißenfelser Straße 42\n04229 Leipzig", ref="MP-260916-33", signatures=(("Sven Riedel", "Vertrieb Muldenpack"),))
    mail(0, "22_bestellung_kantine_450g.eml", "kunde1", ["ott"], "2026-09-17T09:20:00+02:00", "Bestellung KB-261017 / 2.400 Beutel Hausmischung 450 g", """Guten Morgen Frau Ott,

wir bestellen wie am Dienstag besprochen 2.400 Beutel Hausmischung, jeweils 450 g ganze Bohnen, zu 9,20 EUR netto pro Beutel. Gesamt 22.080,00 EUR netto zuzüglich der für den Kaffee geltenden Umsatzsteuer. Anlieferung am 1. Oktober bis 12 Uhr in unserem Lager Nonnenstraße 14. Die Muster vom Mai und das grüne Etikett bleiben maßgeblich.

Wir planen damit unsere Ausgabe ab 2. Oktober. Bitte melden Sie sich bis kommenden Dienstag, falls die Menge nicht vollständig kommt. Ein Wechsel auf 250-g-Beutel passt nicht in unsere Ausgabe und müsste vorher mit mir abgestimmt werden. Eine Teilmenge nehmen wir nicht ungefragt an.

Freundliche Grüße""")
    mail(0, "23_bestellung_hotels_450g.eml", "kunde2", ["ott"], "2026-09-18T14:08:00+02:00", "Abruf MH-26-104 / 1.800 Beutel 450 g zum 6. Oktober", """Sehr geehrte Frau Ott,

hiermit rufen wir 1.800 Beutel Elsterbogen Frühstücksröstung zu jeweils 450 g ab. Vereinbart sind 8,90 EUR netto je Beutel, insgesamt 16.020,00 EUR netto zuzüglich Umsatzsteuer. Bitte liefern Sie am 06.10.2026 zwischen 7 und 10 Uhr an unser Zentrallager Brühl 38. Unser Rahmenpreis gilt unverändert.

Die Ware wird am selben Tag auf drei Häuser verteilt. Nach dem letzten beschädigten Karton möchten wir diesmal vor Versand die Kartonanzahl wissen. Eine andere Beutelgröße können wir wegen der bereits gedruckten Zimmerkarten nicht verwenden. Frau Seidel erwähnte Ihre neue Anlage. Ich hoffe, die Einweisung ist inzwischen abgeschlossen; geben Sie mir bitte eine kurze Mengenbestätigung.

Mit freundlichen Grüßen""")
    path_for(0, "24_produktion_uebertragung.xlsx", "Produktion PX12, Übertragung des Produktionsbuchs")
    path_for(0, "25_zahlungen_px12.xlsx", "Rechnungen und Zahlungen PX12")
    machine_screen()


def mainz():
    mail(1, "01_kraemer_an_reuter_21-09.eml", "kraemer", ["reuter"], "2026-09-21T08:45:00+02:00", "Torben Hesse / Unterlagen vor unserer Antwort", """Sehr geehrte Frau Reuter,

bitte begleiten Sie uns als Unternehmensführung in der Angelegenheit Torben Hesse. Wir haben 17 Beschäftigte, keinen Betriebsrat und haben den Vertrieb verkleinert. Herr Hesse hat die Kündigung inzwischen erhalten. Er widerspricht dem Enddatum, fordert den Bonus für 2025 und nennt 22 offene Urlaubstage. Wir haben darauf außer zur Laptoprückgabe noch nicht geantwortet.

Anbei sind der Arbeitsvertrag und der später unterschriebene Nachtrag. Derya hat beim Erstellen des Briefs die frühere Vertragsdatei geöffnet. Ich habe den Brief am 11. September unterschrieben. Die Vorabmail ging am Samstag heraus, der Bote war am Dienstag beim Haus. Bitte sprechen Sie mit uns, bevor Sie Herrn Hesse oder seiner Seite etwas anbieten. Wir wären zu einer geordneten Trennung bereit, haben aber keinen Betrag beschlossen.

Den Bonus haben wir noch nicht ausgezahlt. Herr Berg und Derya kommen bei drei Aufträgen nicht zum selben Umsatzanteil wie Herr Hesse. Zusätzlich steht der Kundenrabatt von 1.800 EUR im Raum. Herr Berg hatte ihm vorher ausdrücklich erlaubt, Zahlen für das Budget aus dem CRM mitzunehmen. Dass dafür ein privater USB-Stick benutzt wurde, habe ich erst bei der Rückgabe erfahren. Bitte klären Sie mit uns, welche Unterlagen Sie für die Antwort und ein mögliches Gespräch noch benötigen.

Mit freundlichen Grüßen""", attachments=("02_arbeitsvertrag_2019_unterschrieben.pdf", "03_nachtrag_2023_unterschrieben.pdf", "09_kuendigung_original.pdf", "11_einwurfnachweis_15-09.pdf"))
    pdf(1, "02_arbeitsvertrag_2019_unterschrieben.pdf", MUHL, "Arbeitsvertrag Torben Hesse", "Mainz, 15.02.2019", [
        "Zwischen der Mühlentor Sensortechnik GmbH, vertreten durch Marlene Krämer, und Herrn Torben Hesse, Wallaustraße 37, 55118 Mainz, wird folgender Arbeitsvertrag geschlossen.",
        "1 Tätigkeit und Beginn",
        "Herr Hesse wird ab 01.04.2019 unbefristet als Vertriebsmitarbeiter am Standort Mainz beschäftigt. Er betreut Bestandskunden, akquiriert Industriekunden und dokumentiert Angebote und Aufträge im CRM. Der Arbeitgeber kann ihm gleichwertige Aufgaben zuweisen, soweit diese seinen Kenntnissen entsprechen. Kundenbesuche und gelegentliche Dienstreisen gehören zur Tätigkeit.",
        "2 Arbeitszeit und Vergütung",
        "Die regelmäßige Arbeitszeit beträgt 40 Stunden in einer Fünftagewoche von Montag bis Freitag. Beginn und Ende werden mit der Vertriebsleitung abgestimmt. Angeordnete Mehrarbeit wird durch Freizeit ausgeglichen. Das Monatsgehalt beträgt 4.300,00 EUR brutto und ist am letzten Bankarbeitstag des Monats zu zahlen. Zusätzlich erhält Herr Hesse einen umsatzabhängigen Bonus von zwei Prozent des ihm zugeordneten Nettoumsatzes.",
        "Für den Bonus zählt der im Kalenderjahr ausgelieferte und fakturierte Umsatz ohne Umsatzsteuer, Fracht und durchlaufende Fremdleistungen. Bei gemeinsamer Kundenbetreuung wird die Zuordnung im CRM vor der Abrechnung durch die Vertriebsleitung mitgeteilt. Gutschriften vermindern den zugehörigen Umsatz; spätere Gutschriften werden in der nächsten Abrechnung unter Nennung des Ursprungsauftrags berücksichtigt. Die Abrechnung erfolgt nach Abschluss des Kalenderjahres, die Zahlung bis 30. April des Folgejahres. Vorschüsse bedürfen einer gesonderten Vereinbarung.",
        "3 Urlaub und Verhinderung",
        "Der Jahresurlaub beträgt 30 Arbeitstage bei der vereinbarten Fünftagewoche. Urlaub wird vor Antritt beantragt und von der Vertriebsleitung genehmigt. Eine Erkrankung und deren voraussichtliche Dauer sind unverzüglich mitzuteilen. Dauert sie länger als drei Kalendertage, ist eine ärztliche Bescheinigung nach den jeweils geltenden Vorgaben vorzulegen.",
        "4 Arbeitsmittel und Vertraulichkeit",
        "Der Arbeitgeber stellt Laptop, Telefon und dienstliche Zugänge. Kundendaten dürfen nur für betriebliche Aufgaben verwendet werden. Eine Speicherung auf privaten Datenträgern setzt die vorherige Zustimmung der Vertriebsleitung voraus. Nach Beendigung der Nutzung sind betriebliche Dateien an den Arbeitgeber zurückzugeben und private Kopien in Abstimmung mit der IT zu löschen. Private Dateien dürfen nicht in gemeinsame Kundenordner eingestellt werden. Geschäfts- und Kundengeheimnisse sind auch nach Ende des Arbeitsverhältnisses vertraulich zu behandeln.",
        "5 Kündigung und Nebenabreden",
        "Die ersten sechs Monate gelten als Probezeit; währenddessen kann mit zwei Wochen Frist gekündigt werden. Danach kann jede Partei mit vier Wochen zum Monatsende kündigen. Gesetzliche Verlängerungen einer vom Arbeitgeber einzuhaltenden Frist bleiben unberührt. Kündigungen bedürfen der Schriftform. Änderungen dieses Vertrags sollen schriftlich festgehalten werden; individuelle Vereinbarungen bleiben möglich. Ein Tarifvertrag wird nicht in Bezug genommen. Jede Partei erhält eine von beiden Seiten unterschriebene Ausfertigung."
    ], ref="Personalnummer 019", signatures=(("Marlene Krämer", "Für Mühlentor Sensortechnik GmbH, 15.02.2019"), ("Torben Hesse", "Arbeitnehmer, 15.02.2019")))
    pdf(1, "03_nachtrag_2023_unterschrieben.pdf", MUHL, "Nachtrag zum Arbeitsvertrag vom 15. Februar 2019", "Mainz, 20.11.2023", [
        "Mühlentor Sensortechnik GmbH, Rheinallee 88, 55120 Mainz, vertreten durch Marlene Krämer, und Torben Hesse, Wallaustraße 37, 55118 Mainz, vereinbaren mit Wirkung zum 01.01.2024 die folgenden Änderungen.",
        "1 Funktion und Monatsgehalt",
        "Herr Hesse übernimmt die Funktion Senior Vertrieb Industriekunden. Der Arbeitsort Mainz und die regelmäßige Wochenarbeitszeit von 40 Stunden bleiben unverändert. Das feste Monatsgehalt beträgt ab 01.01.2024 4.800,00 EUR brutto. Die Bonusregelung in Ziffer 2 des Arbeitsvertrags bleibt einschließlich des Satzes von zwei Prozent und der Abrechnung zum 30. April des Folgejahres bestehen. Ein pauschaler Bonusabschlag wird nicht vereinbart.",
        "2 Kündigungsfrist",
        "Nach Ablauf der bereits beendeten Probezeit beträgt die Kündigungsfrist für beide Parteien drei Monate zum Ende eines Kalendervierteljahres. Diese Vereinbarung ersetzt die vertragliche Frist von vier Wochen zum Monatsende in Ziffer 5 des Arbeitsvertrags. Längere zwingende gesetzliche Fristen bleiben unberührt. Die Kündigung bedarf weiterhin der Schriftform.",
        "3 Fortgeltung und Ausfertigungen",
        "Die übrigen Bestimmungen des Arbeitsvertrags vom 15.02.2019, insbesondere der Jahresurlaub von 30 Arbeitstagen, bleiben unverändert. Dieser Nachtrag umfasst eine Seite und wird in zwei gleichlautenden Ausfertigungen unterschrieben. Eine Ausfertigung verbleibt bei Herrn Hesse, eine in der Personalakte."
    ], ref="Personalnummer 019 / Nachtrag 1", signatures=(("Marlene Krämer", "Geschäftsführerin, 20.11.2023"), ("Torben Hesse", "Arbeitnehmer, 20.11.2023")))
    pdf(1, "04_geschaeftsfuehrung_vertriebsumbau.pdf", MUHL, "Entscheidung zur Vertriebsorganisation", "31.08.2026", [
        "1 Auftragslage und Kapazität",
        "Der Auftragseingang Januar bis August beträgt 1,86 Mio. EUR netto gegenüber 2,41 Mio. EUR im Vorjahreszeitraum. Serienkunden haben Wiederholabrufe um vier bis sechs Monate verschoben. Der Außendienst Industriekunden besteht neben Vertriebsleiter Nils Berg aus Torben Hesse, Nora Seifert, David Roth, Mina Scholz und Emil Weber. Weitere elf Beschäftigte arbeiten in Entwicklung, Fertigung, Versand, Verwaltung und IT.",
        "2 Künftige Aufgabenverteilung",
        "Ab November soll eine Stelle im Außendienst entfallen. Herr Hesses bisherige Serienkunden sollen von Frau Seifert betreut werden. Herr Roth übernimmt Neuanfragen zu Standardmodulen. Herr Berg führt die beiden Rahmenvertragsverhandlungen selbst. Die Geschäftsführung hat Herrn Hesses Stelle dafür benannt, weil seine bisherigen Gebiete an die beiden benachbarten Regionen anschließen. Ob Schulungsaufwand und laufende Reisezeiten dadurch steigen, wird aus den Kalendern noch zusammengetragen.",
        "Für die telefonisch betreuten kleinen Bestellungen bleibt Herr Weber zuständig. Frau Scholz bearbeitet Integrationsprojekte mit Softwareanpassung. Die in der Augustplanung genannte zusätzliche Stelle Anwendungssupport wird vorerst nicht ausgeschrieben. Frau Scholz soll zunächst an zwei Tagen pro Woche diesen Support übernehmen. Eine dauerhafte Besetzung ist noch nicht beschlossen.",
        "3 Umsetzung",
        "Derya Engel bereitet das Schreiben an Herrn Hesse vor. Nils Berg soll die Übernahmen mit den betroffenen Beschäftigten besprechen und die Kundenzuordnungen dokumentieren. Die Bestandsliste vom 31. August weist 17 Beschäftigte aus. Eine Arbeitnehmervertretung besteht nicht. Die Geschäftsführerin ist in dieser Zahl nicht enthalten. Es wird keine weitere Kündigung beschlossen."
    ], ref="GF 2026-08-31", signatures=(("Marlene Krämer", "Geschäftsführerin"),))
    personnel = [
        ['Personalnr','Name','Geburtsdatum','Eintritt','Wochenstunden','Tätigkeit','Unterhalt laut Personalbogen','Schwerbehinderung laut Personalbogen'],
        ['019','Torben Hesse','1984-06-18','2019-04-01',40,'Senior Vertrieb Industriekunden','verheiratet; 2 Kinder, geboren 2014 und 2017','keine Angabe; Bogen 2020'],
        ['007','Nils Berg','1976-10-12','2012-08-01',40,'Vertriebsleitung','verheiratet; 1 Kind','nein'],
        ['021','Nora Seifert','1989-03-24','2020-02-01',40,'Außendienst Serienkunden','ledig; 1 Kind','nein'],
        ['024','David Roth','1994-12-07','2022-06-01',40,'Außendienst Standardmodule','ledig; 0 Kinder','nein'],
        ['016','Mina Scholz','1987-08-15','2017-09-01',32,'Vertrieb Integration und Software','verheiratet; 2 Kinder','nein'],
        ['028','Emil Weber','1998-01-30','2024-03-01',30,'Telefonvertrieb Kleinbestellungen','ledig; 0 Kinder','nein'],
        ['004','Derya Engel','1979-05-11','2011-05-01',35,'Personal und Finanzen','verheiratet; 1 Kind','nein'],
        ['009','Sven Kühn','1982-09-05','2013-02-01',40,'Elektronikentwicklung','verheiratet; 2 Kinder','nein'],
        ['011','Lea Fink','1985-02-16','2014-10-01',40,'Firmwareentwicklung','ledig; 0 Kinder','nein'],
        ['014','Philipp Kern','1990-07-21','2016-04-01',40,'Prüffeld','ledig; 0 Kinder','nein'],
        ['017','Rita Maas','1970-04-19','2018-01-01',30,'Fertigung','verheiratet; 0 Kinder','nein'],
        ['018','Cem Aydin','1988-11-26','2018-09-01',40,'Fertigung','verheiratet; 2 Kinder','nein'],
        ['020','Karla Behr','1992-06-04','2019-10-01',40,'Fertigung','ledig; 0 Kinder','nein'],
        ['023','Jens Wolf','1981-12-14','2021-03-01',40,'Versand','verheiratet; 1 Kind','nein'],
        ['025','Anja Grün','1995-09-02','2022-11-01',30,'Einkauf','ledig; 0 Kinder','nein'],
        ['026','Oskar Seitz','1997-03-18','2023-04-01',40,'IT und Geräteverwaltung','ledig; 0 Kinder','nein'],
        ['027','Mara Koch','1993-10-27','2023-09-01',25,'Auftragsabwicklung','verheiratet; 1 Kind','nein'],
    ]
    csv_file(1, "05_personalbestand_31-08-2026.csv", "Personalbestand am 31. August 2026", personnel)
    mail(1, "06_berg_kundenzuordnung.eml", "berg", ["kraemer", "engel"], "2026-09-03T17:12:00+02:00", "Verteilung Industriekunden ab November", """Hallo Marlene, hallo Derya,

Nora kann die Serienkunden von Torben übernehmen, braucht dafür aber Zugriff auf dessen Gesprächsnotizen. David kann die Standardangebote übernehmen; für die Rahmenverträge möchte ich selbst mit den Kunden sprechen. Das betrifft zusammen etwa 60 laufende Kontakte. Die Liste mit der endgültigen Zuordnung ist noch nicht im CRM freigegeben.

Torben hat vor zwei Jahren auch Integrationsprojekte betreut. Mina macht inzwischen die Anpassungen mit der Entwicklung und erklärt den Kunden die Schnittstellen. Ob Torben nach einer Einweisung wieder in diesen Teil könnte, haben wir nicht besprochen. Der Anwendungssupport soll zunächst intern aufgefangen werden. Ein Stellenangebot habe ich weder geschrieben noch veröffentlicht.

Nora fragte, ob sie eine Gehaltserhöhung für das größere Gebiet bekommt. Das ist offen. Für die zusätzlichen Fahrten im November fehlen mir noch die Hotel- und Kundentermine. Derya, bitte nimm nicht die alten Gebietsumsätze als Reisezeitnachweis.

Viele Grüße""")
    txt(1, "07_chat_crm_zahlen_10-09.txt", "Chat Nils Berg und Torben Hesse am 10. September", """Mühlentor Bürochat
Direktnachrichten Nils Berg / Torben Hesse
Export durch Torben Hesse am 18.09.2026, 18:22 Uhr
Zeitzone Europe/Berlin

10.09.2026 15:41 Nils Berg: Ich brauche morgen die Umsätze der Serienkunden für den Budgettermin. Kannst du die Zahlen aus dem CRM ziehen?
10.09.2026 15:43 Torben Hesse: Bin auf dem Weg zum letzten Kunden. Ich mache die Aufstellung heute Abend daheim, dann kann ich die offenen Angebote danebenlegen.
10.09.2026 15:44 Nils Berg: Ja, nimm die Zahlen aus dem CRM mit. Name, Ansprechpartner, Umsatz 2025 und laufendes Angebot reichen. Die Feldservicenotizen brauche ich nicht.
10.09.2026 15:45 Torben Hesse: Im Hotel klappt die Verbindung schlecht. Ich nehme den Export offline mit und schicke dir morgen nur die Tabelle mit den Summen.
10.09.2026 15:47 Nils Berg: Passt. Bitte vor zehn, ich muss um elf mit Marlene rein.
11.09.2026 09:18 Torben Hesse: Tabelle ist im Teamordner Budget. Bei Bergfeld ist die Rechnung nach der Rücknahme noch nicht richtig zugeordnet. Ich habe die ursprüngliche Zahl stehen lassen.
11.09.2026 09:22 Nils Berg: Danke, ich nehme die Gesamtsumme erst einmal so mit. Den Einzelposten kläre ich mit Derya.
""")
    docx(1, "08_gespraech_seifert_18-09.docx", MUHL, "Gespräch mit Nora Seifert", "18.09.2026, 14:10 bis 14:35 Uhr | Aufgeschrieben von Derya Engel", [
        "Nora kam wegen der geplanten Kundenübernahme in mein Büro. Sie sagte, Nils habe ihr Anfang September erklärt, dass Torben bald nicht mehr im Außendienst sei. Sie habe angenommen, er würde den telefonischen Anwendungssupport übernehmen. Von einer Kündigung habe sie erst am 16. September durch Torben erfahren.",
        "Nach Noras Erinnerung waren Torben und sie im Frühjahr bei zwei Kunden gemeinsam. Bei Berghaus hatte Torben die technische Vorbesprechung geführt und Nora die Folgetermine übernommen. Für die Bonuszuordnung habe es keine Besprechung mit beiden gegeben. Sie habe die Änderung des CRM-Anteils auf ihrem Bildschirm gesehen, aber nicht selbst eingegeben. Ob Torben darüber eine Nachricht erhielt, wusste sie nicht.",
        "Nora sagte außerdem, Torben sei in der Woche vom 7. September morgens mehrfach im Kundengespräch gewesen. Sie habe am 9. September gegen 11 Uhr mit ihm zu einem Angebot telefoniert. Ob er danach frei hatte, könne sie nicht sagen. Im Teamkalender sei diese Woche als Urlaub gelb gewesen. Sie habe die Farbe nicht geändert.",
        "Am Ende fragte Nora, wer die Kunden anruft, solange Torben freigestellt ist. Ich habe ihr gesagt, dass Nils die Verteilung noch schriftlich bestätigt. Das Gespräch dauerte rund 25 Minuten. Nora hat diesen Vermerk am 18. September um 16:02 Uhr gelesen und geschrieben: So erinnere ich es auch; beim gemeinsamen Kunden war es Berghaus, nicht Bergfeld. Diese Korrektur ist oben eingearbeitet."
    ], signer="Derya Engel\nPersonal und Finanzen")
    pdf(1, "09_kuendigung_original.pdf", MUHL, "Kündigung Ihres Arbeitsverhältnisses", "Mainz, 11.09.2026", [
        "Sehr geehrter Herr Hesse,",
        "hiermit kündigen wir das mit Ihnen seit dem 01.04.2019 bestehende Arbeitsverhältnis ordentlich zum 30.11.2026, hilfsweise zum nächstzulässigen Termin. Die Entscheidung steht im Zusammenhang mit der Neuordnung unseres Vertriebs und dem Wegfall Ihrer bisherigen Außendienststelle.",
        "Die Übergabe Ihrer Kunden und der dienstlichen Arbeitsmittel stimmen Sie bitte mit Herrn Berg und Herrn Seitz ab. Über eine Freistellung und die noch offenen Urlaubszeiten erhalten Sie eine gesonderte Nachricht. Mit diesem Schreiben wird kein Urlaubszeitraum festgelegt. Bitte setzen Sie die laufende Dokumentation bis zu einer abweichenden Mitteilung fort.",
        "Ein Arbeitszeugnis erstellen wir Ihnen gesondert. Ihre bis zum Ende des Arbeitsverhältnisses abzurechnenden Bezüge werden über die übliche Entgeltabrechnung abgerechnet. Bitte wenden Sie sich wegen der Meldung als arbeitsuchend rechtzeitig an die zuständige Agentur für Arbeit.",
        "Mit freundlichen Grüßen"
    ], recipient=HESSE, ref="Personalnummer 019", signatures=(("Marlene Krämer", "Geschäftsführerin"),))
    mail(1, "10_vorabmail_12-09.eml", "engel", ["hesse_work", "hesse"], "2026-09-12T09:14:00+02:00", "Persönlich / Schreiben vom 11. September vorab", """Sehr geehrter Herr Hesse,

Frau Krämer hat gestern das beigefügte Schreiben unterschrieben. Sie erhalten es hier vorab als Datei. Das Papieroriginal wird Ihnen gesondert überbracht. Ich habe den Botendienst für Dienstag beauftragt, weil unser Büro heute nicht besetzt ist. Bitte nutzen Sie für Rückfragen zunächst diese Adresse.

Herr Berg ist am Montag ab 13 Uhr erreichbar und bespricht dann die Übergaben mit Ihnen. Die Kundentermine für nächste Woche sind im Kalender noch nicht geändert. Bitte sagen Sie nicht selbst sämtliche Termine ab, bevor Herr Berg die Vertretung benannt hat.

Die Nachricht ersetzt die angekündigte Übergabe des Papiers nicht. Den Laptop benötigt Herr Seitz erst nach Absprache; ein Rückgabetermin steht heute noch nicht fest.

Mit freundlichen Grüßen""", attachments=("09_kuendigung_original.pdf",))
    pdf(1, "11_einwurfnachweis_15-09.pdf", "Rheinboten Kurierdienst e.K.\nInhaber Arne Klotz, Mombacher Straße 68, 55122 Mainz\nTelefon 06131 672 81 0 | disposition@rheinboten-kurier.de", "Zustellprotokoll RB-260915-018", "15.09.2026", [
        "Auftraggeberin: Mühlentor Sensortechnik GmbH, Rheinallee 88, 55120 Mainz. Empfänger: Torben Hesse, Wallaustraße 37, 55118 Mainz. Zustellerin: Eva Lorenz. Auftrag übernommen um 08:05 Uhr bei Derya Engel.",
        "Frau Engel zeigte mir das einseitige, mit blauer Tinte von Marlene Krämer unterschriebene Schreiben mit dem Betreff Kündigung Ihres Arbeitsverhältnisses und Datum 11.09.2026. Ich las Empfänger und Datum, verglich es mit dem beigefügten Auftragsdurchschlag und legte es selbst in den Umschlag. Den verschlossenen Umschlag trug ich ohne Zwischenübergabe zur Zustelladresse.",
        "Ankunft 09:37 Uhr. Das Haus besitzt außen links vom Eingang sechs beschriftete Briefkästen. Der dritte von oben trägt Torben Hesse. Um 09:42 Uhr habe ich den Umschlag vollständig durch dessen Briefschlitz eingeworfen. Er fiel in das Fach; nichts ragte heraus. Ich habe nicht geklingelt und niemanden angetroffen. Die Hausnummer 37 war über der Tür lesbar. Ein Foto wurde nicht aufgenommen.",
        "Rückmeldung an die Disposition um 09:45 Uhr. Die Zeiten habe ich unmittelbar in der Tourenliste notiert und dieses Protokoll nach Rückkehr um 11:10 Uhr unterschrieben."
    ], ref="RB-260915-018", signatures=(("Eva Lorenz", "Zustellerin, Mainz, 15.09.2026, 11:10 Uhr"),))
    mail(1, "12_hesse_widerspruch_16-09.eml", "hesse", ["kraemer", "engel"], "2026-09-16T07:52:00+02:00", "Ihr Brief / Enddatum, Bonus und Urlaub", """Sehr geehrte Frau Krämer,

das Original Ihres Schreibens lag gestern, am 15. September, in meinem Briefkasten, als ich gegen 18 Uhr heimkam. Die Vorabmail vom Samstag hatte ich am Sonntag gelesen. Mit dem Ausscheiden zum 30. November bin ich nicht einverstanden. Auf meinem unterschriebenen Nachtrag steht drei Monate zum Quartalsende. Eine Erklärung, warum meine Kunden nun Nora und David machen sollen, habe ich noch nicht erhalten.

Außerdem fehlen weiterhin 14.200 EUR Bonus für 2025. Das sind zwei Prozent von 710.000 EUR. Die drei nachträglich gekürzten Zuordnungen habe ich nie bestätigt. Den Bergfeld-Rabatt hatte Nils mit dem Kunden besprochen; ich habe dazu die Gutschrift nicht ausgelöst. Bitte zahlen Sie den Bonus aus und schicken Sie mir die vollständige Abrechnung.

Ich habe noch 22 Urlaubstage. Anfang September sollte ich trotz eingetragenen Urlaubs die Kundentermine übernehmen; Nils sagte, die Tage kämen zurück. Bitte streichen Sie die fünf Tage nicht einfach. Ich möchte auch ein ordentliches Zeugnis und die Gründe für die Kundenzuordnung wissen. Für eine vernünftige Übergabe bin ich erreichbar.

Mit freundlichen Grüßen""")
    pdf(1, "13_entgeltabrechnung_august_2026.pdf", MUHL, "Entgeltabrechnung August 2026", "31.08.2026", [
        "Arbeitnehmer: Torben Hesse, Wallaustraße 37, 55118 Mainz. Personalnummer 019. Eintritt 01.04.2019. Abrechnungszeitraum 01.08. bis 31.08.2026. Steuerklasse 4, kein Kirchensteuerabzug. Krankenversicherung gesetzlich. Regelmäßige Wochenarbeitszeit 40 Stunden.",
        "Das feste Monatsgehalt ist vollständig abgerechnet. Die Zeile Bonus Vorjahr beträgt im August 0,00 EUR. Seit Januar wurde in keiner Entgeltabrechnung 2026 ein Bonus für 2025 ausgezahlt. Die separate Umsatzaufstellung ist nicht Bestandteil der Augustabrechnung. Es ist kein Betrag für den Kundenrabatt Bergfeld von diesem Monatsentgelt abgezogen.",
        "Überweisung des Auszahlungsbetrags am 31.08.2026 auf das hinterlegte Gehaltskonto mit Endziffern 4872. Rückfragen zur Abrechnung nimmt Derya Engel entgegen. Die Beträge unter Jahreswerte umfassen Januar bis August einschließlich dieses Monats."
    ], ref="PN 019 / 2026-08", table=[['Lohnart / Abzug', 'Monat EUR', 'Jahr EUR'], ['Festgehalt brutto', '4.800,00', '38.400,00'], ['Bonus Vorjahr', '0,00', '0,00'], ['Lohnsteuer', '-702,21', '-5.617,68'], ['Krankenversicherung Arbeitnehmer', '-420,00', '-3.360,00'], ['Rentenversicherung Arbeitnehmer', '-446,40', '-3.571,20'], ['Arbeitslosenversicherung Arbeitnehmer', '-62,40', '-499,20'], ['Pflegeversicherung Arbeitnehmer', '-74,40', '-595,20'], ['Solidaritätszuschlag / Kirchensteuer', '0,00', '0,00'], ['Auszahlungsbetrag', '3.094,59', '24.756,72']], widths=[290,100,100])
    customers = ['Altenhof Gerätebau','Berghaus Anlagen','Cramer Prüftechnik','Dornbach Systeme','Eichen Messgeräte','Feldmann Labor','Gartenbach Pumpen','Heller Maschinen','Ilmen Werkzeuge','Jäger Automation','Kastner Anlagen','Lichtenberg Technik','Mahlow Geräte','Norden Laborbau','Ostheim Pumpen','Peters Prüfstände','Rabenfeld Industrie','Steinbach Geräte','Ufer Messtechnik','Waldheim Automation']
    amounts = [20000,25000,30000,35000,40000,45000,50000,30000,25000,40000,35000,45000,30000,40000,35000,30000,50000,30000,40000,35000]
    assert sum(amounts) == 710000
    reductions = {1:15000, 8:20000, 16:25000}
    data = [['Auftrag','Kunde','Rechnung','Leistungsdatum','Netto_EUR','Hesse_Anteil_EUR','Abrechnung_Anteil_EUR','Änderung_am','Vermerk_Buchhaltung']]
    for i, (customer, amount) in enumerate(zip(customers, amounts)):
        month = 1 + i // 2
        data.append([f'V25-{101+i}',customer,f'RE25-{301+i}',f'2025-{month:02}-{12 if i%2==0 else 24}',amount,amount,amount-reductions.get(i,0),'2026-04-21' if i in reductions else '',{1:'15.000 an Seifert; CRM-Freigabe Berg',8:'20.000 an Roth; Änderung ohne Empfangsbestätigung',16:'25.000 an Berg; gemeinsamer Rahmenvertrag'}.get(i,'Zuordnung unverändert')])
    csv_file(1, "14_crm_bonus_2025_rohdaten.csv", "CRM-Auszug Bonus 2025 mit Buchhaltungszuordnung", data)
    path_for(1, "15_bonusabgleich_2025.xlsx", "Bonusabrechnung 2025, Derya Engel")
    mail(1, "16_engel_bonus_verrechnung.eml", "engel", ["kraemer", "berg"], "2026-04-28T11:24:00+02:00", "Bonus Hesse 2025 / Freigabe fehlt", """Hallo Marlene, hallo Nils,

die 20 Rechnungen ergeben 710.000 EUR netto. Torben setzt alles sich selbst zu und kommt auf 14.200 EUR. Nach den drei von Nils am 21. April geänderten Anteilen verbleiben in unserer Liste 650.000 EUR und damit 13.000 EUR. Ich habe keine von Torben bestätigte Zuordnungsnachricht gefunden. Die Auslieferungen und Rechnungen selbst sind nicht streitig; die Kürzung um zusammen 60.000 EUR betrifft nur die interne Zurechnung.

Marlene möchte daneben die 1.800 EUR aus der Bergfeld-Gutschrift verrechnen. Das ergäbe nach unserem Ansatz 11.200 EUR. Es geht um den Nettobetrag; der Kunde hat 2.142 EUR einschließlich Umsatzsteuer gutgeschrieben bekommen. Die Gutschrift betrifft einen Auftrag aus 2026 und ist in den 20 Vorjahresrechnungen nicht enthalten.

Ich habe weder die 13.000 noch die 11.200 EUR in den Zahlungslauf genommen. Bitte sagt mir bis morgen, was freigegeben wird. Eine Erklärung an Torben, mit welcher eigenen Forderung wir verrechnen wollen, habe ich nicht verschickt. Nils, schicke mir bitte auch deine Freigabe des Rabatts oder die Nachricht, aus der hervorgeht, dass Torben ohne dich entschieden hat.

Viele Grüße""")
    pdf(1, "17_gutschrift_bergfeld_2026.pdf", MUHL, "Gutschrift GS26-0410", "10.04.2026", [
        "Sehr geehrte Frau Huber,", "zur Rechnung RE26-178 vom 25.03.2026 über Auftrag V26-044 schreiben wir Ihnen den nachträglich vereinbarten Preisnachlass gut. Gegenstand sind zwölf Sensorsätze MT-40. Die ursprüngliche Rechnung über 18.000,00 EUR netto wird um zehn Prozent vermindert. Die Ware verbleibt bei Ihnen; eine Rücknahme erfolgt nicht.",
        "Der Nachlass wurde am 08.04.2026 telefonisch mit unserem Vertrieb besprochen. In unserem Vorgang ist Torben Hesse als Bearbeiter und Nils Berg als Vertriebsleitung hinterlegt. Eine gesonderte schriftliche Freigabe ist diesem Beleg nicht beigefügt. Die Gutschrift wird mit Ihrer nächsten fälligen Zahlung verrechnet.",
        "Umsatzsteuer-ID DE319482716. Bitte nennen Sie bei Rückfragen die Nummer GS26-0410 und den Ursprungsauftrag V26-044. Dieser Beleg betrifft ausschließlich die im März 2026 ausgelieferte Ware.", "Mit freundlichen Grüßen\nDerya Engel\nBuchhaltung"
    ], recipient="Bergfeld Messsysteme GmbH\nFrau Alina Huber\nSiemensstraße 28\n64289 Darmstadt", ref="GS26-0410", table=[['Position', 'EUR'], ['Nachlass netto', '1.800,00'], ['Umsatzsteuer 19 Prozent', '342,00'], ['Gutschrift brutto', '2.142,00']], widths=[360,130])
    path_for(1, "18_urlaubskalender_2026.xlsx", "Urlaubskalender und Buchungen Torben Hesse 2026")
    pdf(1, "19_urlaubskarte_hesse_2026.pdf", MUHL, "Urlaubskarte Torben Hesse 2026", "Stand 16.09.2026", [
        "Personalnummer 019. Die Urlaubskarte wird von Derya Engel geführt; Herr Hesse trägt die gewünschten Zeiträume ein, Nils Berg bestätigt die Freigabe. Die Arbeitstage beziehen sich auf Montag bis Freitag. Jahresurlaub 30 Tage. Übertrag aus 2025: vier Tage. Der Übertrag wurde am 12.12.2025 wegen der verschobenen Kundenabnahme von Nils Berg bis 31.03.2026 bestätigt.",
        "Für den 2. und 5. Januar wurden zwei übertragene Tage verwendet; zwei weitere übertragene Tage wurden nach Abstimmung vom 20.03.2026 auf den Maiurlaub übertragen. Derya Engel hat die Verlängerung im Kalender am 20.03. eingetragen. Am 16.09. meldet Herr Hesse, die Septemberwoche sei wegen der Kundentermine zurückzubuchen. Nils Berg hat auf dieser Karte hierzu noch nichts gegengezeichnet. Die Genehmigungen unten betreffen die ursprünglichen Anträge."
    ], ref="Urlaub / PN 019", table=[['Antrag am', 'Zeitraum', 'Tage', 'Beantragt', 'Genehmigt'], ['10.12.2025', '02.01. und 05.01.2026', '2', 'T. Hesse', 'N. Berg, 12.12.2025'], ['09.02.2026', '04.05. bis 08.05.2026', '5', 'T. Hesse', 'N. Berg, 10.02.2026'], ['04.05.2026', '20.07. bis 24.07.2026', '5', 'T. Hesse', 'N. Berg, 05.05.2026'], ['11.06.2026', '07.09. bis 11.09.2026', '5', 'T. Hesse', 'N. Berg, 12.06.2026']], widths=[75,145,35,85,150], signatures=(("Torben Hesse", "Anträge wie oben eingetragen"), ("Nils Berg", "Freigaben zu den jeweils genannten Daten; Verlängerung Übertrag am 20.03.2026")))
    mail(1, "20_berg_urlaub_kundentermine.eml", "berg", ["hesse_work"], "2026-09-07T07:36:00+02:00", "Diese Woche / Kunden brauchen Rückmeldung", """Guten Morgen Torben,

ich weiß, dass die Woche als Urlaub eingetragen ist. Nora ist krank und David schafft die technischen Rückfragen nicht. Kannst du heute Altenhof und morgen Berghaus übernehmen? Am Mittwoch wäre das Gespräch mit Nora zur Übergabe hilfreich. Für Donnerstag brauche ich noch die CRM-Zahlen und am Freitag den Budgetstand bis zehn.

Bitte trag die tatsächlichen Termine in deinen Kalender ein. Wir bekommen die freien Tage wieder hin. Ich sage Derya Bescheid, wenn feststeht, was du gemacht hast. Du musst nicht den ganzen Tag im Büro sitzen; die Kunden sollen ihre Antworten bekommen. Mir ist wichtig, dass nichts liegen bleibt.

Gruß""", cc=())
    pdf(1, "21_freistellung_16-09.pdf", MUHL, "Freistellung und Übergabe", "Mainz, 16.09.2026", [
        "Sehr geehrter Herr Hesse,", "ab dem 17.09.2026 stellen wir Sie bis auf Weiteres unter Fortzahlung des festen Monatsgehalts von der täglichen Arbeitsleistung frei. Die Freistellung ist widerruflich. Herr Berg darf Sie für konkret abgestimmte Übergabegespräche kontaktieren. Eine tägliche Bereitschaft im Büro ist nicht vorgesehen.",
        "Urlaub wird mit dieser Erklärung nicht festgelegt oder verrechnet. Ihren Einwand zu den fünf Septembertagen besprechen wir gesondert. Bitte reichen Sie uns die Kundentermine aus dieser Woche ein, soweit sie noch nicht im Teamkalender stehen. Auch der von Ihnen angesprochene Bonus wird getrennt behandelt.",
        "Laptop, Netzteil, Diensttelefon und Zutrittskarte geben Sie bitte am 17.09.2026 um 10 Uhr bei Herrn Seitz ab. Private Dateien können Sie in seiner Anwesenheit sichern. Dienstliche Dateien sollen auf den Unternehmenslaufwerken verbleiben. Über privat gespeicherte Kundendateien bitten wir um konkrete Auskunft zu Dateiname, Speicherort und vorhandenen Kopien. Bitte löschen Sie vor der gemeinsamen Abstimmung keine betrieblichen Dateien.", "Mit freundlichen Grüßen"
    ], recipient=HESSE, ref="Personalnummer 019", signatures=(("Marlene Krämer", "Geschäftsführerin"),))
    pdf(1, "22_rueckgabeprotokoll_17-09.pdf", MUHL, "Rückgabeprotokoll Arbeitsmittel", "17.09.2026, 10:02 bis 10:38 Uhr", [
        "Übergebender: Torben Hesse. Entgegennehmender: Oskar Seitz, IT. Ort: Besprechungsraum Rheinallee 88. Herr Hesse meldete sich am Laptop an. Der Ordner Budget lag auf dem Teamlaufwerk und war von Herrn Berg lesbar. Der lokale Export kunden_20260910.csv blieb auf dem Laptop erhalten.",
        "Herr Hesse erklärte, er habe den Export zusätzlich auf einem privaten USB-Stick, auf dem auch Familienfotos seien. Den Stick habe er heute nicht dabei. Herr Seitz hat den Inhalt des privaten Datenträgers nicht gesehen und keine Kopie davon erstellt. Nach Hesses Angabe umfasst die CRM-Datei 412 Kundenkontakte mit Namen, E-Mail, Telefon und Umsatzzahlen. Eine Übermittlung an Dritte verneinte er.",
        "Der Laptop wird im IT-Schrank verwahrt und nicht neu aufgesetzt. Herr Hesse bittet um einen Termin zum Trennen privater und betrieblicher Dateien auf dem Stick. Herr Seitz sagt hierfür einen Vorschlag für die nächste Woche zu. Mit der Unterschrift wird der Empfang der unten aufgeführten Gegenstände und die Wiedergabe der Erklärungen bestätigt."
    ], ref="IT-RG-260917-19", table=[['Arbeitsmittel', 'Kennung', 'Zustand / Übergabe'], ['Laptop mit Netzteil', 'MT-L019 / Seriennr. 7M019628', 'Gehäuse ohne sichtbaren Schaden; startet'], ['Diensttelefon', 'MT-T019', 'Gerät und Ladekabel übergeben'], ['Zutrittskarte', 'ZK-019', 'Empfangen und gesperrt'], ['Privater USB-Stick', 'Hesses Angabe: silber, 32 GB', 'Nicht vorgelegt; kein Empfang']], widths=[135,165,190], signatures=(("Torben Hesse", "Übergebender"), ("Oskar Seitz", "IT / Empfänger")))
    event_screen(1, "23_it_ereignisse_l019.png", "Geräteereignisse MT-L019 und CRM-Export", [
        ['Zeit_CEST','System','Benutzer','Ereignis','Objekt','Wert','Protokoll_ID'],
        ['2026-09-10 15:59:12','CRM','t.hesse','Anmeldung','Sitzung CRM-9951','erfolgreich','EV-401'],
        ['2026-09-10 16:02:38','CRM','t.hesse','Export gestartet','Kunden mit Umsatz 2025','412 Datensätze','EV-402'],
        ['2026-09-10 16:02:42','MT-L019','t.hesse','Datei erstellt','Downloads/kunden_20260910.csv','188416 Byte','EV-403'],
        ['2026-09-10 16:04:05','MT-L019','t.hesse','USB verbunden','USB-Massenspeicher 32 GB','Gerätekennung 82A119','EV-404'],
        ['2026-09-10 16:04:17','MT-L019','t.hesse','Datei geöffnet','Downloads/kunden_20260910.csv','Tabellenanwendung','EV-405'],
        ['2026-09-10 16:08:44','MT-L019','t.hesse','USB getrennt','Gerätekennung 82A119','ordnungsgemäß','EV-406'],
        ['2026-09-11 09:17:28','Teamlaufwerk','t.hesse','Datei gespeichert','Budget/serienkunden_2027.xlsx','28416 Byte','EV-407'],
        ['2026-09-11 09:21:07','Teamlaufwerk','n.berg','Datei geöffnet','Budget/serienkunden_2027.xlsx','Lesezugriff','EV-408'],
        ['2026-09-17 10:15:21','MT-L019','o.seitz','Bestand aufgenommen','Downloads/kunden_20260910.csv','188416 Byte','EV-409'],
        ['2026-09-17 10:34:02','Verzeichnisdienst','o.seitz','Zugang gesperrt','t.hesse','Web und Fernzugriff','EV-410'],
        ['2026-09-17 11:08:00','IT-Ausgabe','o.seitz','Exportvermerk','USB-Dateikopiervorgänge','nicht protokolliert; keine Zielpfadliste vorhanden','EV-411'],
    ])
    mail(1, "24_datenauskunft_hesse_19-09.eml", "hesse", ["engel", "kraemer"], "2026-09-19T10:18:00+02:00", "Auskunft über meine personenbezogenen Daten", """Sehr geehrte Frau Engel,

bitte erteilen Sie mir Auskunft nach Artikel 15 DSGVO zu den über mich verarbeiteten Daten und senden Sie mir eine Kopie. Mir geht es insbesondere um meine Personalakte, die Änderung meiner Umsatzzuordnung für 2025, die Kommunikation über die Auswahl meiner Stelle und die jetzt erhobenen Geräteprotokolle. Bitte nennen Sie auch Zwecke, Empfänger und die vorgesehene Speicherdauer.

Die Gehaltsabrechnungen habe ich bereits. Ich möchte aber wissen, wer die drei CRM-Zuordnungen geändert hat und ob es dazu E-Mails gibt, in denen ich genannt werde. Bitte beziehen Sie Nachrichten von Herrn Berg an die Geschäftsführung ein. Eine pauschale Liste der Systeme würde mir nicht helfen.

Sie können die Antwort über einen geschützten Abruf bereitstellen und mir den Zugang an diese private Adresse schicken. Ich bitte nicht um fremde Personalakten. Sollte eine Eingrenzung nötig sein, sagen Sie mir bitte konkret, welche Information Ihnen fehlt.

Mit freundlichen Grüßen""")
    mail(1, "25_hesse_usb_rueckgabe.eml", "hesse", ["engel"], "2026-09-18T18:06:00+02:00", "Termin mit Herrn Seitz / privater USB-Stick", """Sehr geehrte Frau Engel,

wie gestern gesagt, ist der Stick mein eigener. Darauf sind auch Fotos meiner Kinder und die Steuerunterlagen meiner Frau. Ich möchte deshalb nicht den ganzen Datenträger unbeaufsichtigt abgeben. Den Kundenexport habe ich für Herrn Berg zu Hause geöffnet. Ich habe ihn nicht an einen anderen Arbeitgeber oder einen Kunden verschickt.

Herr Berg hatte am 10. September geschrieben, ich solle die Zahlen aus dem CRM mitnehmen. Den Chat habe ich noch. Ob damit auch mein privater Stick gemeint war, haben wir nicht ausdrücklich besprochen. Der Export heißt kunden_20260910.csv. Ich weiß nicht mehr, ob mein privater Rechner beim Öffnen automatisch eine weitere lokale Kopie angelegt hat. In einer privaten Cloud habe ich ihn nicht bewusst gespeichert.

Ich kann am Dienstag, 22. September, um 14 Uhr kommen und den Stick zusammen mit Herrn Seitz ansehen. Bitte bestätigen Sie den Termin und sagen Sie, ob ich den privaten Rechner ebenfalls mitbringen soll. Bis dahin verändere ich die Datei nicht.

Mit freundlichen Grüßen""")
    docx(1, "26_personalgespraech_kraemer_engel.docx", MUHL, "Besprechung Krämer und Engel", "21.09.2026, 11:00 bis 11:35 Uhr | Notizen Derya Engel", [
        "1 Vertragsmappe und Versand",
        "Ich habe Frau Krämer beide unterschriebenen Unterlagen gezeigt. Der Nachtrag war im Ordner Gehaltsänderung abgelegt, nicht hinter dem Arbeitsvertrag. Für das Schreiben vom 11. September hatte ich die Datei Vertrag_Hesse_2019 geöffnet. Den 30. November habe ich eingetragen, weil wir mit zwei Übergabemonaten gerechnet hatten. Eine schriftliche Berechnung dazu gibt es nicht. Die Vorabmail ging von mir am 12. September heraus. Das Botenprotokoll kam am 15. September um 12:07 Uhr per Scan zurück.",
        "2 Personalstand und Gespräche",
        "Die Liste vom 31. August enthält alle 17 Beschäftigten. Es gibt keine Auszubildenden, keine Leiharbeitskräfte und keinen Betriebsrat. Frau Krämer selbst ist nicht mitgezählt. Torbens Angaben zu Familie und möglichen besonderen Umständen stammen aus dem Personalbogen 2020; eine aktuelle Abfrage haben wir nicht gemacht. Herr Berg hat bislang nur mit Nora über zusätzliche Kunden gesprochen. Mit Torben gab es vor dem Schreiben kein Gespräch über den Anwendungssupport.",
        "3 Urlaub und Vergütung",
        "Der Kalender bucht 17 Tage ab. Torben zählt nur die zwölf Tage im Januar, Mai und Juli. Für September liegen Nils' Mail und zwei Kundenbesprechungen im Teamkalender vor; vollständige Tagesnachweise habe ich nicht. Den Urlaubseintrag habe ich noch nicht zurückgenommen. Der Bonus ist auch im September noch nicht ausgezahlt. Von Nils fehlen mir weiterhin die Mitteilungen über die drei Umsatzänderungen und die Nachricht zum Bergfeld-Rabatt.",
        "4 Geräte und nächste Woche",
        "Oskar ist am 22. September ab 14 Uhr im Haus. Torbens Terminvorschlag ist noch unbeantwortet. Der Laptop bleibt verschlossen im IT-Schrank. Wir haben die privaten Datenträger nicht gesehen. Die Datenauskunft vom 19. September wurde von mir im Personalpostfach abgelegt; eine Antwort oder ein Export wurde noch nicht verschickt. Frau Krämer möchte vor einer inhaltlichen Antwort mit Frau Reuter sprechen."
    ], signer="Derya Engel\nPersonal und Finanzen")


def readmes():
    descriptions = [
        ("Elsterbogen Rösterei Leipzig", "Unterlagen der Geschäftsführung zum Auftrag einer Beutelverschließmaschine PX12, ihrer Einrichtung und der offenen Restrechnung. Aktenstand: 22. September 2026.", "Elsterbogen Rösterei GmbH, vertreten durch Friederike Ott."),
        ("Mühlentor Sensortechnik Mainz", "Unterlagen der Unternehmensführung zur Trennung von Torben Hesse, zur Vergütung und zur Rückgabe betrieblicher Daten. Aktenstand: 21. September 2026. Die Erklärungen des Arbeitnehmers sind als eigene Quellen enthalten.", "Mühlentor Sensortechnik GmbH, vertreten durch Marlene Krämer."),
    ]
    for slug, (title, intro, role) in zip(SLUGS, descriptions):
        folder = ROOT / 'testakten' / slug
        source_files = sorted(FILES[slug])
        assert len(source_files) == 26
        rows = '\n'.join(f'| `{name}` | {caption} |' for name, caption in source_files)
        notice = '\n'.join('> ' + line if line else '>' for line in NOTICE.splitlines())
        text = f"""# {title}

## 1 Aktenstand

{intro}

Mandantin: {role} Zugeordnetes Plugin: `anwaltschaft-generell`. Die Akte umfasst 26 native Quellen. Excel-Aufstellungen stehen neben den Ursprungsbelegen und ersetzen sie nicht.

<!-- BEGIN gesamt-pdf-section (autogen) -->
## 2 Downloads

{notice}

| Fassung | Download |
| --- | --- |
| Gesamt-PDF | [Lesefassung](gesamt-pdf/{slug}_gesamt.pdf) |
| Akten-ZIP | [Originaldateien](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{slug}.zip) |
| Einzel-PDF-ZIP | [Einzelunterlagen](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{slug}-einzelpdfs.zip) |

Die Archivlinks verweisen auf den veröffentlichten Release. Dieser Aktenstand ist für Version 445.1.0 vorgesehen.
<!-- END gesamt-pdf-section (autogen) -->

## 3 Unterlagen

| Datei | Dokument |
| --- | --- |
{rows}

## 4 Herstellung

Die beiden fallbezogenen Generatoren unter `scripts/build-anwaltschaft-leipzig-mainz-akten.py` und `scripts/build-anwaltschaft-leipzig-mainz-workbooks.mjs` erzeugen die Originaldateien. Die bestehenden zentralen Exporte erzeugen Lesefassung und flache Archive. Zwischenstände und visuelle Prüfdateien entstehen ausschließlich unter `/tmp/anwaltschaft-leipzig-mainz-445-1-0`. Die Kontaktangaben sind ausschließlich für diese Unterlagen angelegt; es werden keine Nachrichten versandt und keine Kontakte hergestellt.
"""
        (folder / 'README.md').write_text(text, encoding='utf-8')
        qa = QA_ROOT / slug
        qa.mkdir(parents=True, exist_ok=True)
        data = {}
        for source in folder.glob('*.csv'):
            with source.open(encoding='utf-8-sig', newline='') as stream:
                data[source.name] = list(csv.reader(stream, delimiter=';'))
        (qa / 'tabellendaten.json').write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')


def module(filename, name):
    spec = importlib.util.spec_from_file_location(name, ROOT / 'scripts' / filename)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def packages():
    combined = module('build-testakte-gesamt-pdf.py', 'akten_gesamt')
    originals = module('build-testakten-release-zips.py', 'akten_originale')
    individual = module('build-testakten-einzelpdf-zips.py', 'akten_einzel')
    for slug in SLUGS:
        folder = ROOT / 'testakten' / slug
        out = QA_ROOT / slug / 'releases'
        out.mkdir(parents=True, exist_ok=True)
        print(combined.build_gesamt_pdf(folder))
        print(originals.build_single(folder, out))
        print(individual.build_single(folder, out))


def visual_qa():
    import shutil
    import subprocess
    import zipfile
    from pypdf import PdfReader
    from PIL import Image, ImageDraw
    individual = module('build-testakten-einzelpdf-zips.py', 'akten_einzel_qa')
    bundled = Path(sys.executable).resolve().parents[2] / 'bin' / 'override' / 'pdftoppm'
    renderer = str(bundled) if bundled.exists() else shutil.which('pdftoppm')
    if not renderer:
        raise RuntimeError('Für die Sichtprüfung wird pdftoppm benötigt.')
    for slug in SLUGS:
        folder = ROOT / 'testakten' / slug
        out = QA_ROOT / slug / 'einzel'
        out.mkdir(parents=True, exist_ok=True)
        archive, count = individual.build_single(folder, out)
        assert count == 26, (slug, count)
        pages = []
        with zipfile.ZipFile(archive) as z:
            for name in sorted(z.namelist()):
                if not name.endswith('.pdf'):
                    continue
                data = z.read(name)
                (out / name).write_bytes(data)
                prefix = out / f'{Path(name).stem}-seite'
                for old in out.glob(f'{prefix.name}-*.png'):
                    old.unlink()
                subprocess.run([renderer, '-r', '108', '-png', str(out / name), str(prefix)], check=True, capture_output=True)
                rendered = sorted(out.glob(f'{prefix.name}-*.png'))
                assert len(rendered) == len(PdfReader(out / name).pages), name
                pages.extend(rendered)
        for start in range(0, len(pages), 8):
            montage = Image.new('RGB', (1400, 1040), 'white')
            draw = ImageDraw.Draw(montage)
            for offset in range(min(8, len(pages)-start)):
                im = Image.open(pages[start+offset])
                im.thumbnail((340, 490))
                x, y = (offset % 4)*350, (offset//4)*520
                montage.paste(im, (x, y+22))
                draw.text((x+8, y+4), pages[start+offset].stem[:45], fill='black')
            montage.save(out / f'kontakt-{start+1:03}.png')
        print(f'{slug}: {len(pages)} Seiten aus {count} Quellen gerendert')


def verify():
    from email import policy
    from email.parser import BytesParser
    from openpyxl import load_workbook
    from PIL import Image, ImageStat
    from pypdf import PdfReader
    from testakte_zip_common import working_dump_flat_pairs
    quality = module('validate-testakten-dokumentqualitaet.py', 'akten_dokumentqualitaet')
    for slug in SLUGS:
        folder = ROOT / 'testakten' / slug
        pairs = working_dump_flat_pairs(folder, include_gesamt_pdf=False)
        assert len(pairs) == 26, (slug, len(pairs))
        assert {p.suffix for p, _ in pairs} >= {'.docx', '.pdf', '.eml', '.xlsx', '.csv', '.png'}
        assert not (folder / '.qa').exists()
        report = []
        for p, arc in pairs:
            assert '/' not in arc and '\\' not in arc
            text = quality.export_text(p) if p.suffix in {'.docx', '.pdf', '.eml', '.csv', '.txt'} else ''
            if p.suffix in {'.docx', '.pdf', '.eml', '.txt'}:
                assert len(text.strip()) >= 600, (p, len(text.strip()))
                assert not quality.language_prose_errors(text, p), (p, quality.language_prose_errors(text, p))
            assert not re.search(r'testakte|fiktiv|platzhalter|formathinweis|lösungsmatrix|musterlösung|§|\.(?:example|test|invalid)\b', text, re.I), p
            if p.suffix == '.eml':
                assert not quality.eml_quality_errors(p), (p, quality.eml_quality_errors(p))
                msg = BytesParser(policy=policy.default).parsebytes(p.read_bytes())
                for attached in str(msg.get('X-Attachments', '')).split(','):
                    if attached.strip():
                        assert (folder / attached.strip()).exists(), (p, attached)
            elif p.suffix == '.pdf':
                assert quality.pdf_is_a4(p), p
            elif p.suffix == '.docx':
                assert quality.is_a4(Document(p)), p
            elif p.suffix == '.csv':
                with p.open(encoding='utf-8-sig', newline='') as f:
                    rows = list(csv.reader(f, delimiter=';'))
                assert len(rows) >= 5 and all(len(row) == len(rows[0]) for row in rows), p
            elif p.suffix == '.png':
                with Image.open(p) as im:
                    assert im.width >= 900 and ImageStat.Stat(im.convert('L')).stddev[0] > 15, p
            elif p.suffix == '.xlsx':
                values = load_workbook(p, data_only=True)
                formulas = load_workbook(p, data_only=False)
                count = 0
                for sheet in formulas:
                    assert sheet.freeze_panes, p
                    for row in sheet:
                        for cell in row:
                            if cell.data_type == 'f':
                                count += 1
                                cached = values[sheet.title][cell.coordinate]
                                assert cached.data_type != 'e' and cached.value is not None, (p, cell.coordinate)
                assert count > 0, p
                values.close()
                formulas.close()
            report.append({'datei': p.name, 'textzeichen': len(text.strip())})
        qa = QA_ROOT / slug
        qa.mkdir(parents=True, exist_ok=True)
        (qa / 'quellenpruefung.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(f'{slug}: 26 native Quellen, Mindestlängen, Header, Formeln und Bildbeleg geprüft')
    def cell(case, file, coord):
        w = load_workbook(ROOT / 'testakten' / SLUGS[case] / file, data_only=True)
        value = w.active[coord].value
        w.close()
        return value
    assert cell(0, '24_produktion_uebertragung.xlsx', 'D25') == 990
    assert cell(0, '24_produktion_uebertragung.xlsx', 'E25') == 65
    assert cell(0, '24_produktion_uebertragung.xlsx', 'G25') == 5.5
    assert cell(0, '25_zahlungen_px12.xlsx', 'E6') == 42840
    assert cell(0, '25_zahlungen_px12.xlsx', 'E8') == 28560
    assert cell(0, '25_zahlungen_px12.xlsx', 'B12') == 14280
    assert cell(1, '15_bonusabgleich_2025.xlsx', 'E27') == 14200
    assert cell(1, '15_bonusabgleich_2025.xlsx', 'F27') == 13000
    assert cell(1, '15_bonusabgleich_2025.xlsx', 'F32') == 11200
    assert cell(1, '18_urlaubskalender_2026.xlsx', 'C27') == 17
    assert cell(1, '18_urlaubskalender_2026.xlsx', 'D28') == 22
    print('Fallbezogene Summen und Differenzen stimmen.')


def run():
    p = argparse.ArgumentParser()
    p.add_argument('--pakete', action='store_true')
    p.add_argument('--qa', action='store_true')
    p.add_argument('--pruefen', action='store_true')
    args = p.parse_args()
    if args.pruefen:
        verify()
    elif args.pakete:
        packages()
    elif args.qa:
        visual_qa()
    else:
        init_fonts()
        screen_fonts(12)
        leipzig()
        mainz()
        readmes()
        for slug in SLUGS:
            print(f'{slug}: {len(FILES[slug])} Quellen registriert; Excel-Dateien mit dem Tabellen-Generator bauen.')


if __name__ == '__main__':
    run()
