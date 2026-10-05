#!/usr/bin/env python3
"""Prüft Quellenintegrität und Rechnungen der drei kommunalen Haftpflichtakten.

Rechtliche Ergebnisse werden getrennt anhand der fallbezogenen Evaluationsrubrik
beurteilt; dieser Test behauptet keine automatische juristische Freigabe.
"""
from email import policy
from email.parser import BytesParser
import io
import os
from pathlib import Path
import re
import unittest
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from kommunale_haftpflicht_akten_daten import CASES
from testakte_disclaimer import NOTICE_DE, NOTICE_EN, NOTICE_FILENAME, NOTICE_TEXT

ROOT = Path(__file__).resolve().parents[1]


def directory(case):
    return ROOT / 'testakten' / case['slug']


def inventory(case):
    return {d['file'] for d in case['documents']} | set(case['xlsx'])


class MunicipalCasesTest(unittest.TestCase):
    def test_exact_native_inventory(self):
        for case in CASES:
            actual = {p.name for p in directory(case).iterdir()
                      if p.is_file() and p.name not in {'README.md', 'rubric.yaml'}}
            self.assertEqual(inventory(case), actual, case['slug'])
            self.assertEqual(len(actual), 12)
            self.assertTrue({'.docx', '.eml', '.txt'} <= {Path(n).suffix for n in actual})

    def test_email_headers_body_and_actual_mime_attachments(self):
        for case in CASES:
            for item in case['documents']:
                if not item['file'].endswith('.eml'):
                    continue
                message = BytesParser(policy=policy.default).parsebytes((directory(case)/item['file']).read_bytes())
                for header in ['From', 'To', 'Subject', 'Date', 'Message-ID']:
                    self.assertTrue(message[header], (item['file'], header))
                    self.assertFalse(message[header].defects)
                body = message.get_body(preferencelist=('plain',)).get_content()
                self.assertEqual(body.replace('\r\n', '\n').strip(), item['body'])
                parts = list(message.iter_attachments())
                attachments = {p.get_filename(): p.get_payload(decode=True) for p in parts}
                self.assertEqual(len(parts), len(attachments))
                self.assertEqual(set(attachments), set(case['attachments'].get(item['file'], [])))
                for name, data in attachments.items():
                    self.assertEqual(data, (directory(case)/name).read_bytes(), name)

    def test_word_files_preserve_complete_authored_paragraphs(self):
        for case in CASES:
            for item in case['documents']:
                if item['file'].endswith('.docx'):
                    text = '\n'.join(p.text for p in Document(directory(case)/item['file']).paragraphs)
                    for paragraph in re.split(r'\n\s*\n', item['body']):
                        self.assertIn(paragraph.removeprefix('## '), text, item['file'])

    def test_workbook_formulas_have_valid_cached_results(self):
        expected = {
            'akha-wuerzburg-rohrbruch': ('Kostenliste', {'D7': 2400, 'D8': 1800, 'D9': 7500, 'D10': 3400, 'D12': 15100}),
            'akha-wuerzburg-betriebsfahrzeug': ('Forderungen', {'D7': 145000, 'D8': 1650000, 'D9': 45000, 'D11': 1840000}),
        }
        for case in CASES:
            for name in case['xlsx']:
                path = directory(case)/name
                formulas = load_workbook(path, data_only=False)
                cached = load_workbook(path, data_only=True)
                sheet, cells = expected[case['slug']]
                for coordinate, value in cells.items():
                    self.assertEqual(formulas[sheet][coordinate].data_type, 'f', coordinate)
                    self.assertEqual(cached[sheet][coordinate].value, value, coordinate)
                for s in cached:
                    for row in s:
                        for cell in row:
                            self.assertNotEqual(cell.data_type, 'e', (name, cell.coordinate))
                if case['slug'].endswith('betriebsfahrzeug'):
                    self.assertEqual(formulas[sheet]['D11'].value.upper(), '=SUM(D7:D9)')
                    self.assertEqual(cached[sheet]['C16'].value, 1200000)
                    self.assertEqual(cached[sheet]['C17'].value, 2100000)
                    self.assertIn('nicht', cached[sheet]['A19'].value)

    def test_small_core_set_and_plugin_navigation(self):
        for case in CASES:
            readme = (directory(case)/'README.md').read_text()
            self.assertLessEqual(len(case['core']), 6)
            for name in case['core']:
                self.assertIn(f']({name})', readme)
                self.assertIn(name, inventory(case))
            self.assertIn('../../kommunale-haftpflicht/README.md', readme)
            self.assertIn('reserved-example-contacts', readme)
            self.assertIn(NOTICE_DE, readme)
            self.assertIn(NOTICE_EN, readme)

    def test_combined_pdfs_have_complete_document_bookmarks(self):
        for case in CASES:
            path = directory(case)/'gesamt-pdf'/f"{case['slug']}_gesamt.pdf"
            reader = PdfReader(path)
            text = '\n'.join(p.extract_text() for p in reader.pages)
            self.assertNotIn(NOTICE_DE, text)
            self.assertNotIn(NOTICE_EN, text)
            self.assertGreaterEqual(len(reader.pages), 12)
            outline = str(reader.outline)
            for name in inventory(case):
                self.assertIn(Path(name).stem, outline, name)

    @unittest.skipUnless(os.environ.get('KOMMUNALE_HAFTPFLICHT_ZIPS'), 'KOMMUNALE_HAFTPFLICHT_ZIPS für Archivprüfung setzen')
    def test_flat_archives_preserve_sources_and_one_pdf_per_original(self):
        for case in CASES:
            from testakte_zip_common import working_dump_archive_pairs
            from testakte_einzelpdf_common import document_arcname_pairs
            native_pairs = {arc:p for p,arc in working_dump_archive_pairs(directory(case), include_gesamt_pdf=False)}
            native = set(native_pairs)
            for suffix in ('', '-einzelpdfs'):
                path = Path(os.environ['KOMMUNALE_HAFTPFLICHT_ZIPS'])/f"testakte-{case['slug']}{suffix}.zip"
                with zipfile.ZipFile(path) as archive:
                    self.assertIsNone(archive.testzip())
                    names = archive.namelist()
                    self.assertEqual(len(names), len(set(names)))
                    self.assertTrue(all('/' not in n for n in names))
                    notice = archive.read(NOTICE_FILENAME).decode('utf-8')
                    self.assertTrue(notice.startswith(NOTICE_TEXT))
                    self.assertIn(NOTICE_EN, notice)
                    if suffix:
                        expected = {arc for p,arc in document_arcname_pairs(directory(case))} | {NOTICE_FILENAME}
                        self.assertEqual(expected, set(names))
                        for name in expected-{NOTICE_FILENAME}:
                            pages = PdfReader(io.BytesIO(archive.read(name))).pages
                            self.assertGreater(len(pages), 0)
                            text = '\n'.join(p.extract_text() for p in pages)
                            self.assertNotIn(NOTICE_DE, text)
                            self.assertNotIn(NOTICE_EN, text)
                    else:
                        combined = f"{case['slug']}_gesamt.pdf"
                        self.assertEqual(native | {NOTICE_FILENAME, combined}, set(names))
                        for name in native:
                            self.assertEqual(archive.read(name), native_pairs[name].read_bytes(), name)
                        self.assertEqual(archive.read(combined), (directory(case)/'gesamt-pdf'/combined).read_bytes())


if __name__ == '__main__':
    unittest.main()
