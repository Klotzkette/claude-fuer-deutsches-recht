#!/usr/bin/env python3
"""Installationsgrenzen und vollständige Seminarakten, keine Modellbewertung."""

import csv
from decimal import Decimal
from email import policy
from email.parser import BytesParser
import hashlib
import io
import json
from pathlib import Path
import re
import unittest
import unicodedata
from urllib.parse import unquote

import yaml
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook
from prompt_profiles import validate_files, hand_curated
from quality_lab import validate_profile
from testakte_einzelpdf_common import document_arcname_pairs
from testakte_zip_common import working_dump_archive_pairs

ROOT = Path(__file__).resolve().parents[1]
CATALOG = json.loads((ROOT / "bauwirtschaft-rundum/seminarfaelle.json").read_text())
EXPECTED = {
    "bauwirtschaft-anfaenger": {"bau-rundum-begehung-einbeck", "bau-rundum-lv-abgleich-detmold", "bau-rundum-bieterfragen-celle", "bau-rundum-behinderung-soest", "bau-rundum-baugrund-verden"},
    "bauwirtschaft-fortgeschrittene": {"bau-rundum-abschlagsrechnung-lemgo", "bau-rundum-vgv-bewerbung-hameln", "bau-rundum-angebotspruefung-goslar", "bau-rundum-bauzeit-claim-minden", "bau-rundum-nachtrag-northeim"},
}
ORIGINAL_COUNTS = {
    "bau-rundum-begehung-einbeck": 14,
    "bau-rundum-lv-abgleich-detmold": 12,
    "bau-rundum-bieterfragen-celle": 15,
    "bau-rundum-behinderung-soest": 12,
    "bau-rundum-baugrund-verden": 12,
    "bau-rundum-abschlagsrechnung-lemgo": 16,
    "bau-rundum-vgv-bewerbung-hameln": 16,
    "bau-rundum-angebotspruefung-goslar": 16,
    "bau-rundum-bauzeit-claim-minden": 16,
    "bau-rundum-nachtrag-northeim": 16,
}


def compact(text):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text).replace("\u00ad", ""))


def source_paragraphs(path):
    if path.suffix == ".docx":
        document = Document(path)
        yield from (p.text for p in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from (p.text for p in cell.paragraphs)
    elif path.suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        body = message.get_body(preferencelist=("plain",))
        if body:
            yield from re.split(r"\n\s*\n", body.get_content())
    elif path.suffix == ".txt":
        yield from re.split(r"\n\s*\n", path.read_text(encoding="utf-8-sig"))
    elif path.suffix == ".pdf":
        yield from (page.extract_text() or "" for page in PdfReader(path).pages)


class SeminarSuite(unittest.TestCase):
    def test_two_plugins_and_exact_ten_announced_cases(self):
        self.assertEqual({p["name"] for p in CATALOG["plugins"]}, set(EXPECTED))
        self.assertFalse((ROOT / "bauwirtschaft-rundum/.claude-plugin/plugin.json").exists())
        market = json.loads((ROOT / ".claude-plugin/marketplace.json").read_text())
        entries = {p["name"]: p for p in market["plugins"]}
        for plugin in CATALOG["plugins"]:
            name = plugin["name"]
            self.assertEqual({c["slug"] for c in plugin["cases"]}, EXPECTED[name])
            self.assertEqual(len(plugin["cases"]), 5)
            self.assertEqual(entries[name]["source"], plugin["source"])
            directory = ROOT / plugin["source"]
            for kind in (".claude-plugin", ".codex-plugin"):
                manifest = json.loads((directory / kind / "plugin.json").read_text())
                self.assertEqual(manifest["name"], name)
                self.assertEqual(manifest["version"], market["version"])
            self.assertEqual(entries[name]["description"], json.loads((directory / ".claude-plugin/plugin.json").read_text())["description"])

    def test_eight_discoverable_skills_and_install_local_references(self):
        for plugin in CATALOG["plugins"]:
            directory = ROOT / plugin["source"]
            skills = list((directory / "skills").glob("*/SKILL.md"))
            self.assertEqual(len(skills), 8)
            for path in skills:
                with self.subTest(skill=path):
                    text = path.read_text()
                    metadata = yaml.safe_load(text.split("---", 2)[1])
                    self.assertEqual(set(metadata), {"name", "description"})
                    self.assertEqual(metadata["name"], path.parent.name)
                    self.assertGreaterEqual(len(metadata["description"]), 80)
                    self.assertLessEqual(len(metadata["description"]), 1024)
                    self.assertLessEqual(len(path.read_bytes()), 7500)
                    self.assertRegex(text, r"(?m)^# \S")
                    for target in re.findall(r"\]\(([^\s)]+)\)", text):
                        if ":" in target or target.startswith("#"):
                            continue
                        dest = (path.parent / unquote(target.split("#")[0])).resolve()
                        self.assertTrue(dest.is_relative_to(directory.resolve()), f"Nicht installiert: {target}")
                        self.assertTrue(dest.is_file(), f"Fehlende Referenz: {target}")

    def test_handwritten_prompts_and_reviewed_profiles(self):
        for plugin in CATALOG["plugins"]:
            name = plugin["name"]
            directory = ROOT / plugin["source"]
            self.assertTrue(hand_curated(name))
            self.assertEqual(validate_files(directory, name, ROOT), [])
            profile = json.loads((ROOT / "quality/evals" / f"{name}.json").read_text())
            validate_profile(profile, name, directory, ROOT)
            for kind, key in (("werkstatt", "workshop_review"), ("schnellstart", "mini_review")):
                data = (directory / f"{name}-{kind}.md").read_bytes()
                self.assertEqual(hashlib.sha256(data).hexdigest(), profile[key]["sha256"])
                self.assertNotIn("§", data.decode())
            self.assertLessEqual(len((directory / f"{name}-schnellstart.md").read_bytes()), 7500)

    def test_navigation_follows_the_two_seminar_levels(self):
        labels = {
            "bauwirtschaft-anfaenger": "Begehung und Planungsabgleich",
            "bauwirtschaft-fortgeschrittene": "Rechnung, Bauzeit und Nachtrag",
        }
        for plugin in CATALOG["plugins"]:
            text = (ROOT / plugin["source"] / "README.md").read_text()
            block = text.split("<!-- BEGIN SKILLS-LOGIC (auto-generated) -->", 1)[1].split("<!-- END SKILLS-LOGIC", 1)[0]
            self.assertIn(labels[plugin["name"]], block)
            self.assertNotIn("Spezialmodule und Schnittstellen", block)
            self.assertEqual(block.count("/SKILL.md)"), 8)

    def test_native_case_inventory_and_pdf_correspondence(self):
        for plugin in CATALOG["plugins"]:
            for item in plugin["cases"]:
                folder = ROOT / "testakten" / item["slug"]
                with self.subTest(case=folder.name):
                    source = {p for p in folder.iterdir() if p.suffix.lower() in {".pdf", ".docx", ".xlsx", ".eml", ".csv", ".txt", ".png", ".jpg"}}
                    self.assertGreaterEqual(len(source), 10)
                    self.assertGreaterEqual(len({p.suffix.lower() for p in source}), 4)
                    pairs = working_dump_archive_pairs(folder, include_gesamt_pdf=False)
                    self.assertEqual({p for p, _ in pairs}, source, "Ein Aktenstück fällt aus dem Export.")
                    self.assertEqual(len(pairs), len({n for _, n in pairs}))
                    pdf_pairs = document_arcname_pairs(folder)
                    self.assertEqual(len(pdf_pairs), len(source))
                    self.assertTrue(all("/" not in n for _, n in pairs + pdf_pairs))
                    pdf = folder / "gesamt-pdf" / f"{folder.name}_gesamt.pdf"
                    self.assertGreaterEqual(len(PdfReader(pdf).pages), len(source))
                    text = "\n".join(p.extract_text() or "" for p in PdfReader(pdf).pages)
                    self.assertNotIn("Diese Testakte wurde", text)
                    readme = (folder / "README.md").read_text()
                    self.assertIn(f"`{plugin['name']}`", readme)
                    self.assertIn("<!-- reserved-example-contacts -->", readme)
                    self.assertIn("This test case file was generated with AI", readme)

    def test_email_headers_actual_attachments_and_csv_shape(self):
        for case_set in EXPECTED.values():
            for slug in case_set:
                folder = ROOT / "testakten" / slug
                for path in folder.glob("*.eml"):
                    with self.subTest(email=path):
                        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                        for header in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version"):
                            self.assertTrue(message[header], header)
                        self.assertIsNotNone(message.get_body(preferencelist=("plain",)))
                        for attachment in message.iter_attachments():
                            filename = attachment.get_filename()
                            if filename and (folder / filename).is_file():
                                self.assertEqual(attachment.get_payload(decode=True), (folder / filename).read_bytes())
                for path in folder.glob("*.csv"):
                    with self.subTest(csv=path):
                        text = path.read_text(encoding="utf-8-sig")
                        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
                        rows = list(csv.reader(io.StringIO(text), dialect))
                        self.assertGreaterEqual(len(rows), 5)
                        self.assertTrue(all(len(row) == len(rows[0]) for row in rows))

    def test_original_text_survives_combined_pdf_export(self):
        for slug in ORIGINAL_COUNTS:
            folder = ROOT / "testakten" / slug
            pdf = PdfReader(folder / "gesamt-pdf" / f"{slug}_gesamt.pdf")
            combined = compact("\n".join(page.extract_text() or "" for page in pdf.pages))
            checked = 0
            for path, _ in working_dump_archive_pairs(folder, include_gesamt_pdf=False):
                for paragraph in source_paragraphs(path):
                    content = compact(paragraph)
                    if len(content) < 80:
                        continue
                    with self.subTest(case=slug, source=path.name, paragraph=paragraph[:60]):
                        self.assertTrue(content in combined, f"Originaltext fehlt im Gesamt-PDF: {paragraph[:160]}")
                    checked += 1
            self.assertGreater(checked, 20, f"{slug}: zu wenige vollständige Originalabsätze geprüft.")

    def test_each_case_gains_at_least_four_independent_documents(self):
        for slug, previous_count in ORIGINAL_COUNTS.items():
            folder = ROOT / "testakten" / slug
            files = [p for p, _ in working_dump_archive_pairs(folder, include_gesamt_pdf=False)]
            additions = [p for p in files if int(p.name.split("_", 1)[0]) > previous_count]
            with self.subTest(case=slug):
                self.assertGreaterEqual(len(files), previous_count + 4)
                self.assertGreaterEqual(len(additions), 4)
                self.assertGreaterEqual(len({p.suffix for p in additions}), 2)
                readme = (folder / "README.md").read_text()
                for path in files:
                    self.assertIn(f"]({path.name})", readme, "Original fehlt im Aktenverzeichnis.")

    def test_lemgo_position_ids_and_actual_payment_chain(self):
        folder = ROOT / "testakten/bau-rundum-abschlagsrechnung-lemgo"
        book = load_workbook(folder / "11_Abschlagsrechnung_3.xlsx", data_only=True)
        sheet = book["WB26098"]
        expected = ["01.010", "02.010", "02.020", "03.010", "03.020", "04.010", "04.020", "05.010"]
        self.assertEqual([sheet.cell(row, 1).value for row in range(6, 14)], expected)
        self.assertTrue(all(sheet.cell(row, 1).data_type == "s" for row in range(6, 14)))
        self.assertAlmostEqual(sheet["F16"].value, 151522.50, places=2)
        self.assertAlmostEqual(sheet["F22"].value, 70657.11, places=2)
        with (folder / "06_Kreditorenbewegungen.csv").open(encoding="utf-8-sig", newline="") as stream:
            payments = sum((Decimal(row["Betrag_EUR"]) for row in csv.DictReader(stream, delimiter=";") if row["Art"] == "Auszahlung"), Decimal(0))
        self.assertEqual(payments, Decimal("100639.08"))
        self.assertEqual(Decimal(str(sheet["B20"].value)) + Decimal(str(sheet["B21"].value)), payments)
        book.close()


if __name__ == "__main__":
    unittest.main()
