#!/usr/bin/env python3
"""Prüft erhaltene Kernakten, Nachträge und das Komponentenrelease 445.34.1.

Ohne --dist werden Quellen und die abgelegten Gesamt-PDFs geprüft. Mit --dist
werden zusätzlich sämtliche nativen Archivbytes und PDF-Inhalte geprüft.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import unicodedata
import unittest
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
import pdfplumber
from pypdf import PdfReader
import yaml

from gesellschafterstreit_nachtrag_daten import CASES
from prompt_profiles import PROMPT_SUFFIXES, validate_files
from testakte_disclaimer import NOTICE_BYTES, pdf_content_errors
from testakte_einzelpdf_common import document_arcname_pairs
from testakte_zip_common import working_dump_archive_pairs

ROOT = Path(__file__).resolve().parents[1]
PLUGIN = ROOT / "gesellschafterstreit"
VERSION = "445.34.1"
ADDENDUM = "Nachtrag_2026-10-08"
DIST = None
# Unabhängiger SHA-256-Bestand der 76 vor dem Update vorhandenen Originale.
# README, Bewertungsdatei und abgeleitetes Gesamt-PDF gehören nicht dazu.
BASELINE_SHA256 = {
    "gesellschafterstreit-zink-und-zunder": {
        "01_Mandatsauftrag_Romy.eml": "6b92556818e6fbcd5277c7ea5f0acc5f320f8b329795c7a99d1ebe4a5752853a",
        "02_Gesellschaft_und_Personen.docx": "bdf1ee49a143fdfd4fa8af988df20634ad4d67a70bfc472424fbe0de1477b843",
        "03_Satzung_Arbeitsabschrift_2023.docx": "65939475971c25c5aec564c1f8ce8790b97f8a4ae470101ac735f7a26197fc3e",
        "04_Bestellung_und_Vertretung.docx": "eadf225863269757a429709ac397bd5b6649d458b45c9b105e2e618144868f9f",
        "05_Beiratsordnung.docx": "fa9905f95c60d45d680d6e7b56ccae60883cf79935577a5e6e5795e7e08a3e16",
        "06_Darlehen_Romy_60000.docx": "119f2ac661d96f60273d93807585d9860028039eb25489eee785bc5ef3642b0b",
        "07_Darlehen_Kunibert_40000.docx": "874e56526905b7504bc8ebde815eca86ede76b909a6e7bae751f7a5f725f23f8",
        "08_Beirat_Anfrage_Darlehen.eml": "50e9df2c543d51a2e7136845a813863a93348f14d5edacac33743d66d3d2f615",
        "09_Beirat_Zwischenantwort.eml": "1a6ae95ab8f7578d4d2e9e82e2b792d5a7d0f0b49946164a17b124a7172a5950",
        "10_Zahlungsnachweis_15_September.docx": "4f268e2951f3f05c4a7975bc86e4fef85768aa2d8b5a3f0098d97dcf95dadbab",
        "11_Steuerbuero_Rueckfragen.eml": "97a4c8d91703f40ec540dbe8bd36c418f8f8793eebf2c111ef22c3fe6e75adfa",
        "12_Finanzgespraech_Kapitalbedarf.docx": "8107378dac2031a465f01838c1239885327689be7f895c6f53b131cd07506427",
        "13_Unverbindliches_Angebot_Hanna.docx": "9b26aca2a509be1c2f4a9fc766fd710015328b0aae70ace4c47159d6c364bf38",
        "14_Hanna_Angebotsversand.eml": "6b37d6ce65dda1b019bf5e94d517f26613ceed520e2c514feac850c03ec5dfc1",
        "15_Romy_Garantien_und_Preis.eml": "85d332dc23bf85997307db6bf37049e6ccfc4753537a8987aee108c7fe716a9b",
        "16_Chat_Verkaeufergruppe.txt": "6c95ca66c4856c6e08fcb48290802bc03213cc9a5e1d1049c286db0e7d17bf32",
        "17_Fensterfuchs_Projektprofil.docx": "c64fa465f7434309c8e92021877d872eabc96d93b2adcd6d8560b09a263aef34",
        "18_Kundenanfrage_Altbogen.eml": "9e6c84703127a873ccf0e813834b4c1a78f2dc636d8485e712dea6507382e9bc",
        "19_CAD_Zugriffsprotokoll.txt": "25d4c952e82b1f5875a0af329925282267314740f962cf37b90bdcfe1ade6ce2",
        "20_IT_Vermerk_und_Rechte.docx": "4da7e99592fe412767ed1d3d71d6ccebdf4c9966f9097fbb2efee527bdcbb352",
        "21_Kunibert_Stellungnahme.eml": "016559d2d16296b8b87b2b7e08557b6b618cbfdc02822388768ceeeef2a3cbde",
        "22_Einladung_09_September.docx": "0ecda189af7bfdc0065c40353517b21cfaffdd14e10279492cde435be260a12d",
        "23_Einladung_Zugang.eml": "036e4c1e2f2b3646d1ade743441f67aac861436c458005d756730653ffaac544",
        "24_Romy_Ergaenzungsantrag.eml": "16d6236f9286e8bbcd7fb9d873a0b1a47aa3e450548fa4874bc7146e87684970",
        "25_Tagesordnung_Nachtrag.docx": "ac47391c68116a881cecd9efb9daf622b3d0f5fe7ab9d3a756ab5b9b428395b9",
        "26_Nachtrag_Zugang.eml": "0e08b48049dd1223e6cc43978df00d166032de35261d26eec5c15f7717aa3b55",
        "27_Protokoll_25_September.docx": "9526db6d9abe7cfa3f3f76e86f1837b80d80eb99f7e57f16ee9e5ba7e9b38756",
        "28_Theklas_Mitschrift.docx": "255b763417cbdbc4af678b14dd1320a08682cd2132987b1241520e03612c6f9c",
        "29_Protokollversand_28_September.eml": "0b2e205556b5ec375715317c21a48b28c517f062a60d75a57ba22b9c8ca58a61",
        "30_Romy_Zugang_und_Widerspruch.eml": "daf8f5395928ce5271c8ee708e4361cc0750c31b539731c4265d8a60908c2104",
        "31_Beirat_Stand_30_September.docx": "5f37e18f892f5e53e702c404f32727f06b2072a52b7e75c85c0b48269781740e",
        "32_Romys_Mandantennotiz.docx": "a19ce8101f03f196200924ca925cfbb1ad784ccbe8144b79c06f8616656a5ea0",
        "33_Genehmigung_Entwurf_ungezeichnet.docx": "f7f61ce4e097ce3ad6f8d2875cbffd5a704c76a0aaacaf0a47e6306beff3ab7b",
        "34_Bueronotiz_Vollzugsstand.docx": "2285ca7a71af6887e4eef88c5b1a6862c3726b3b6746cf781016829d5a1d10e1",
        "35_Rueckfrage_Abfindung.eml": "627b3af996e3f4ac02c3407466d54ddee3b69737791bacfec14f3a2075387931",
        "36_Samir_Hinweis_Ablage.eml": "4ccc47cbc428b1eaa0a17c62d83c17a0b1e699c6c5943ea5b8890389e0db981d",
        "37_Finanzuebersicht.xlsx": "053f320634c1e44732c01e48bcfb88664b39dbc06538386e516f6c925dc4c31e",
        "38_Stimmen_und_Kapital.xlsx": "293a58ec93dd18f8a1efddc270dc288d561ee21a3fa26ba7fa3a63bf42ddd0c2"
    },
    "gesellschafterstreit-klageerwiderung-berlin": {
        "01_Mandat_Rabenstein.eml": "d76b34bcc4b297437346d71138db8929c5b881de8e438e751c5fc5ac24a0d49f",
        "02_Klage_Seidel.docx": "51f78aa75d3142c885948b770260864ce42e56a2c66a371d203ecbffdfdfcd27",
        "03_Gerichtliche_Verfuegung.pdf": "73499e682184245f25315695fa935378644decee060a1bcd99e7230f5f18e2d0",
        "04_Zustellung_Posteingang.txt": "645a11c36c558b668b72520039a175115ccd340e80dcde222d0dbccc60489728",
        "05_Gesellschaftsvertrag_K1.docx": "2e238465c283d85d689ecea441799551e0879e196c37cc6bf97e232bcb827bdb",
        "06a_Gesellschafterliste_K2.docx": "1d1a1de4ab28be5826a61151514e4b4cd6beb6ab87f7a11494a67346a72de9b5",
        "06b_Registerabruf_K3.pdf": "25ee81e2c1ec16740996a75039a7d34e421b510679378ba1e7e15f338dcdfdb8",
        "07_Einladung_K4.docx": "2d9af8860fd1aaf831d45b35d8f122dd9e2cb05a54c6a91d2393579f9d34de30",
        "08_Niederschrift_K5.docx": "b8259f149cb517d88feb33f791d25270a77a358373ee64020b74a64c4148b374",
        "09_Geschaeftsordnung.docx": "72f44ffac09b523fb21e7ce32c9099b2fee392ca61e52d7339c222bdc3509932",
        "10_Rechnung_SBS_K6.pdf": "c018dceef44cd3aa15f97568a64f96bd99bc50896a68dd16e47e7777548e00a9",
        "11_Nachrichten_K7.txt": "b8b935e4ec49c8d9f5c21cd61176575c24340195fe3410a853eb3214283e2a81",
        "12_Beanstandung_K8.docx": "4f44761e7826949eaaaf51c25aadbccb5df704b99141811842afc58ec8781881",
        "13_Antwort_Seidel_K9.docx": "117348cce12e19fc2f15e0947dd3095a41c1bb0542e546c2ada1ff684756d726",
        "14_Wareneingang_Lindenhof.docx": "036f9e8fb6d3912bab8dc9476fafa3443a418c996c6cc84e9ef8cbefd1d5408a",
        "15_Zeuge_Brandt.eml": "6d22e4b0e63b15c199c9fb6258f4697dbcc9b701ea4c85c309eea46f777d2b3e",
        "16_Rechnungseingang.eml": "d717efe8534731d4a76de62e6749ee9c8c21f0e458cde31b1be9879b06aea2c6",
        "17_Bankprotokoll.csv": "1b313697113c47fa063d5002b1fbc3972d68bdfca6af623b700c7a057ad2a71e",
        "18_Telefonvermerk_Heller.docx": "f03a4e20e17deabfb829fb1417fe38ec13cc1a08832ff7de26d60397cd448776",
        "19_Projektkonto.xlsx": "6f662ed66968a1c3a28c19e56012eed43ad39c6c7fbca93ca4cc410bb7d80ac7"
    },
    "gesellschafterstreit-shareholder-agreement-muenchen": {
        "01_Mandat_Hildegard.eml": "9635e2838f5922d1169c92916360620b27508833fb31fbabdec7c6d5deed3ef1",
        "02_Satzung_2021.docx": "cf19923988dbdf8603784284311af8188a5c19ca444b68df48efc9157be0b207",
        "03_Gesellschafterliste.docx": "c015a451a634b1340d939e10e26f40755fd89d6904ca652f65eb0208c92ebe03",
        "04_Angebot_Investor.eml": "dd9e99190c80cdc57eea3f52f7a167d8208a735dfec0b0fc6d59c47e20f6e7ff",
        "05_Eckpunkte_Vogl.docx": "55c144cdfb472a9587a4ade16984663005d1c9a19d6a36f2e24bfd93bb92ef2a",
        "06_Entwurf_Quirin.docx": "618ca99141060c331b7b9c533f7d90bdcb612765e78819ac6110c6c526d78f51",
        "07_Hildegard_Rueckmeldung.eml": "551537ea399491b29b2f2b8e74ec823c422571af8d326d26fd01bc413bdeb38c",
        "08_Chat_Gruender.txt": "e3bda74feb2e4dd127ef368e738d08d532616a74755b76f2521a722f2c86b8ce",
        "09_Geschaeftsfuehrervertrag_Rottmayer.docx": "768edaddfaae6f1b9be1de0407e01ee1a0e42d6db30c8d977fdc7ec5c7f98df0",
        "10_Entwicklungsbestand.docx": "b89ef8a5b14063099f80e7c29241ebec88bc301edaf2ec9b20582d740afd2365",
        "11_Horn_Nutzungsrechte.eml": "e7f8db54ba88085aa8ce72d6cba704bdab303a1941fafc12314d21d70ff61e6e",
        "12_Pruefstand_Bestellung.pdf": "bcd4c9736e6c9d2890d6dc9975a74f509c64fd8902a076054205f254e6833961",
        "13_Bank_Rueckfrage.eml": "809c76a2fa9e9889e6121b80f1b55d03542b1da42f33281782b1dc29968dac70",
        "14_Auftragsbuch.csv": "d902351b6bb5cf560aa96d941897a963eb2ce998b060febb2619d9073ff66e4a",
        "15_Telefonnotiz_Vogl.docx": "a015f0e36f72cc4d0be67308483e2be2eedee2b0683ea0f46f5b1957c0068f84",
        "16_Notariat_Termin.eml": "3359330f17291aadfcce58aed768ffd2c0670c7a703b179837b7c093973ebd10",
        "17_Budgetnotiz_Ottmar.docx": "6efca884b5fc16b7384ec07078f53ba472e2df9666c25fd7206854c6f62c6847",
        "18_Beteiligungsrechnung.xlsx": "6ead802205ac68207a01221e919ad32b4f72e2242db7bd33a17f046c5bc40484"
    }
}


def module(filename):
    spec = importlib.util.spec_from_file_location(filename.replace("-", "_"), ROOT / "scripts" / filename)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


def normalized(text):
    """Nur Unicode-/Umbruchnormalisierung; Zahlen und Satzzeichen bleiben."""
    text = unicodedata.normalize("NFKC", text).replace("\u00ad", "")
    return re.sub(r"\s+", "", text)


def pdf_text(data):
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)


def docx_pdf_body_text(data):
    """Nur geometrisch belegte laufende Rand-Seitenzahlen ausnehmen.

    Eine Zahl muss allein in ihrer Zeile, im äußeren 50-pt-Rand und passend
    zur PDF-Seite stehen. Erst dieselbe Position auf mindestens zwei Seiten
    belegt eine laufende Seitennummer; gewöhnliche Textzahlen bleiben erhalten.
    """
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        candidates = []
        for number, page in enumerate(pdf.pages, 1):
            words = page.extract_words()
            for word in words:
                if word["text"] != str(number):
                    continue
                edge = "top" if word["bottom"] <= 50 else "bottom" if word["top"] >= page.height - 50 else None
                if edge is None:
                    continue
                same_line = [other for other in words if
                             min(word["bottom"], other["bottom"]) - max(word["top"], other["top"]) > .5]
                if len(same_line) != 1:
                    continue
                position = (edge, round(word["x1"], 0),
                            round(word["top"] if edge == "top" else page.height - word["bottom"], 0))
                candidates.append((number - 1, position, word))
        repeated = Counter(position for _, position, _ in candidates)
        if not any(count >= 2 for count in repeated.values()):
            return pdf_text(data)
        texts = []
        for index, page in enumerate(pdf.pages):
            excluded = [word for page_index, position, word in candidates
                        if page_index == index and repeated[position] >= 2]
            def keep(obj):
                if obj.get("object_type") != "char":
                    return True
                return not any(word["x0"] - .1 <= obj["x0"] and obj["x1"] <= word["x1"] + .1
                               and word["top"] - .1 <= obj["top"] and obj["bottom"] <= word["bottom"] + .1
                               for word in excluded)
            texts.append(page.filter(keep).extract_text() or "")
        return "\n".join(texts)


def docx_passages(path):
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as archive:
        tree = ET.fromstring(archive.read("word/document.xml"))
    return ["".join(node.text or "" for node in paragraph.iter(ns + "t"))
            for paragraph in tree.iter(ns + "p")]


def addon_files(case):
    names = [d["file"] for d in case["docs"]]
    names += [str(Path(name).with_suffix(".pdf")) for name in names]
    return names + [mail["file"] for mail in case["emails"]] + [case["xlsx"], case["chat"]]


def source_files(case):
    directory = ROOT / "testakten" / case["slug"]
    return [directory / name for name in BASELINE_SHA256[case["slug"]]] + [
        directory / ADDENDUM / name for name in addon_files(case)
    ]


def release_copies():
    return [PLUGIN / f"gesellschafterstreit-{kind}.{ext}"
            for kind in ("werkstatt", "schnellstart", "hauptproblem") for ext in ("md", "txt")] + [
        ROOT / "testakten" / case["slug"] / "gesamt-pdf" / f"{case['slug']}_gesamt.pdf" for case in CASES
    ] + [ROOT / "docs/handbuecher" / name for name in (
        "gesellschafterstreit-skills-handbuch.pdf", "gesellschafterstreit-werkstatt-lesefassung.pdf")]


def numeric_texts(cell):
    """Darstellungsvarianten gleicher Werte: deutsches/englisches Zahlenformat."""
    value, fmt = cell.value, cell.number_format
    if isinstance(value, (datetime, date)):
        return [value.strftime(pattern) for pattern in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y")]
    if isinstance(value, bool):
        return ["TRUE", "WAHR"] if value else ["FALSE", "FALSCH"]
    if not isinstance(value, (int, float)):
        return [str(value)]
    clean_format = re.sub(r'"[^"]*"|\[[^]]*\]', "", fmt.split(";")[0])
    match = re.search(r"[0#]\.([0#]+)", clean_format)
    precision = len(match.group(1)) if match else 0
    percent = "%" in fmt
    number = value * 100 if percent else value
    english = format(abs(number), f",.{precision}f")
    german = english.translate(str.maketrans(",.", ".,"))
    candidates = [english, german]
    if value < 0:
        candidates = [prefix + text + suffix for text in candidates for prefix, suffix in (("-", ""), ("(", ")"))]
    if percent:
        candidates = [text + "%" for text in candidates]
    elif not match and "#,##" not in clean_format:
        candidates += [format(value, "g")]
    if value == 0 and len(fmt.split(";")) > 2:
        zero_format = fmt.split(";")[2]
        if "–" in zero_format or '"-"' in zero_format:
            candidates += ["–", "-"]
    return candidates


def workbook_pdf_views(data):
    """Extrahiert ganze Zeilen und erkannte Tabellen-/Spaltenansichten."""
    views = [normalized(pdf_text(data))]
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            views.append(normalized(page.extract_text() or ""))
            for table in page.extract_tables():
                for row in table:
                    views.extend(normalized(value) for value in row if value)
            # Office-PDFs haben häufig keine gedruckten Gitterlinien. Häufige
            # Textanfänge liefern dann Spaltengrenzen für umbrochene Zelltexte.
            starts = Counter(round(word["x0"], 1) for word in page.extract_words())
            boundaries = sorted({max(0, x - 1) for x, count in starts.items() if count >= 3} | {0, page.width})
            for left, right in zip(boundaries, boundaries[1:]):
                if right - left > 12:
                    views.append(normalized(page.crop((left, 0, right, page.height)).extract_text() or ""))
    return views


class GesellschafterstreitUpdate(unittest.TestCase):
    def assert_passage(self, passage, text, label):
        passage = re.sub(r"^#{1,6}\s+", "", passage.strip(), flags=re.M)
        if passage:
            self.assertIn(normalized(passage), normalized(text), label)

    def assert_pdf_layout(self, data, label, minimum=9):
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            self.assertTrue(pdf.pages, label)
            for page in pdf.pages:
                self.assertTrue(page.chars, f"Leere Seite: {label}")
                for char in page.chars:
                    if not char["text"].strip():
                        continue
                    self.assertGreaterEqual(char["size"], minimum - .02, f"Unlesbare Schrift: {label}")
                    self.assertGreaterEqual(char["x0"], -.2, label)
                    self.assertLessEqual(char["x1"], page.width + .2, label)
                    self.assertGreaterEqual(char["top"], -.2, label)
                    self.assertLessEqual(char["bottom"], page.height + .2, label)

    def test_docx_pdf_page_numbers_do_not_change_body_numbers(self):
        from reportlab.pdfgen.canvas import Canvas
        for edge in ("top", "bottom"):
            with self.subTest(edge=edge):
                buffer = io.BytesIO()
                canvas = Canvas(buffer, pagesize=(595, 842), invariant=1)
                for page_number in (1, 2):
                    canvas.drawRightString(540, 812 if edge == "top" else 36, str(page_number))
                    canvas.drawString(60, 740, "Vorgang NF-2026-0915; Kaufpreis 25.000 EUR.")
                    canvas.drawString(60, 700, str(page_number))  # Reguläre Zahl im Satzspiegel.
                    canvas.drawString(60, 660, "bedarf keiner" if page_number == 1 else "erneuten Zustimmung.")
                    canvas.drawString(60, 30 if edge == "top" else 810, "2 EUR")  # Keine isolierte Seitenzahl.
                    canvas.showPage()
                canvas.save()
                result = docx_pdf_body_text(buffer.getvalue())
                self.assertEqual(re.findall(r"(?m)^([12])$", result), ["1", "2"])
                self.assertEqual(result.count("NF-2026-0915"), 2)
                self.assertEqual(result.count("25.000 EUR"), 2)
                self.assertEqual(result.count("2 EUR"), 2)
                self.assertIn("bedarf keiner", result)
                self.assertIn("erneuten Zustimmung.", result)

    def test_historical_originals_remain_byte_identical(self):
        self.assertEqual({slug: len(files) for slug, files in BASELINE_SHA256.items()}, {
            "gesellschafterstreit-zink-und-zunder": 38,
            "gesellschafterstreit-klageerwiderung-berlin": 20,
            "gesellschafterstreit-shareholder-agreement-muenchen": 18,
        })
        for slug, files in BASELINE_SHA256.items():
            directory = ROOT / "testakten" / slug
            actual = {p.name for p in directory.iterdir() if p.is_file() and p.name not in {"README.md", "rubric.yaml", ".DS_Store"}}
            self.assertEqual(actual, set(files), slug)
            for name, expected in files.items():
                self.assertEqual(hashlib.sha256((directory / name).read_bytes()).hexdigest(), expected, f"{slug}/{name}")

    def test_exact_addendum_and_export_inventory(self):
        self.assertEqual({case["slug"] for case in CASES}, set(BASELINE_SHA256))
        for case in CASES:
            directory = ROOT / "testakten" / case["slug"]
            expected = addon_files(case)
            self.assertEqual(len(expected), 20)
            self.assertEqual(len(set(expected)), 20)
            self.assertEqual(Counter(Path(name).suffix for name in expected), {".docx": 6, ".pdf": 6, ".eml": 6, ".xlsx": 1, ".txt": 1})
            actual = {p.relative_to(directory / ADDENDUM).as_posix() for p in (directory / ADDENDUM).rglob("*") if p.is_file()}
            self.assertEqual(actual, set(expected) | {"README.md"}, case["slug"])
            selected = {path for path, _ in working_dump_archive_pairs(directory, include_gesamt_pdf=False)}
            self.assertEqual(selected, set(source_files(case)))
            self.assertEqual({path for path, _ in document_arcname_pairs(directory)}, selected)
            for path in selected:
                self.assertGreater(path.stat().st_size, 0, str(path))

    def test_addendum_texts_and_emails_preserve_authored_content(self):
        message_ids = []
        for case in CASES:
            directory = ROOT / "testakten" / case["slug"] / ADDENDUM
            for item in case["docs"]:
                word_text = "\n".join(docx_passages(directory / item["file"]))
                pdf_path = (directory / item["file"]).with_suffix(".pdf")
                data = pdf_path.read_bytes()
                self.assertEqual(pdf_content_errors(data), [], pdf_path.name)
                output_text = pdf_text(data)
                for passage in [item["title"], *re.split(r"\n\s*\n", item["body"])]:
                    self.assert_passage(passage, word_text, item["file"])
                    self.assert_passage(passage, output_text, pdf_path.name)
                self.assert_pdf_layout(data, str(pdf_path))
            self.assertEqual((directory / case["chat"]).read_text(encoding="utf-8").strip(), case["chat_body"].strip())
            for item in case["emails"]:
                message = BytesParser(policy=policy.default).parsebytes((directory / item["file"]).read_bytes())
                for header in ("From", "To", "Subject", "Date", "Message-ID"):
                    self.assertTrue(message[header], (item["file"], header))
                    self.assertFalse(message[header].defects, (item["file"], header))
                self.assertFalse(message.defects)
                self.assertIsNotNone(parsedate_to_datetime(str(message["Date"])).utcoffset())
                message_ids.append(str(message["Message-ID"]))
                self.assertEqual(message.get_body(preferencelist=("plain",)).get_content().replace("\r\n", "\n").strip(), item["body"].strip())
                attachments = list(message.iter_attachments())
                self.assertCountEqual([part.get_filename() for part in attachments], item["attachments"])
                for part in attachments:
                    self.assertEqual(part.get_payload(decode=True), (directory / part.get_filename()).read_bytes(), item["file"])
        self.assertEqual(len(message_ids), len(set(message_ids)))

    def test_workbook_formula_caches_are_present_without_excel_errors(self):
        for case in CASES:
            for path in source_files(case):
                if path.suffix != ".xlsx":
                    continue
                formulas = load_workbook(path, data_only=False)
                cached = load_workbook(path, data_only=True)
                try:
                    count = 0
                    self.assertEqual(formulas.sheetnames, cached.sheetnames)
                    for sheet in formulas:
                        for row in sheet:
                            for cell in row:
                                actual = cached[sheet.title][cell.coordinate]
                                self.assertNotEqual(actual.data_type, "e", f"{path}/{sheet.title}/{cell.coordinate}")
                                if cell.data_type == "f":
                                    count += 1
                                    self.assertIsNotNone(actual.value, f"Formelcache fehlt: {path}/{sheet.title}/{cell.coordinate}")
                    self.assertGreater(count, 0, str(path))
                finally:
                    formulas.close(); cached.close()

    def test_all_email_attachments_match_their_native_sources(self):
        for case in CASES:
            paths = source_files(case)
            by_name = {path.name: path for path in paths}
            self.assertEqual(len(by_name), len(paths), case["slug"])
            for path in paths:
                if path.suffix != ".eml":
                    continue
                message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                for header in ("From", "To", "Subject", "Date", "Message-ID"):
                    self.assertTrue(message[header], (path.name, header))
                    self.assertFalse(message[header].defects, (path.name, header))
                attachments = list(message.iter_attachments())
                names = [part.get_filename() for part in attachments]
                self.assertEqual(len(names), len(set(names)), path.name)
                for part in attachments:
                    self.assertIn(part.get_filename(), by_name, path.name)
                    self.assertEqual(part.get_payload(decode=True), by_name[part.get_filename()].read_bytes(), path.name)

    def test_addendum_calculations_reconcile_with_the_source_documents(self):
        # Unabhängige Belegwerte, keine aus dem Tabellenbuilder übernommene
        # Erwartungsmatrix: Berlin K6/N03; München N03 und alte Bestellung;
        # Zink N01 und die beiden ursprünglichen Darlehensurkunden.
        expected = {
            "gesellschafterstreit-klageerwiderung-berlin": {
                "Rechnung": {"B10": 15462.18, "D10": 18400, "B12": 2937.81,
                             "D12": .01, "D14": 3500, "D21": 1950, "D22": 2320.50},
            },
            "gesellschafterstreit-shareholder-agreement-muenchen": {
                "Kapital": {"D9": 31250, "E6": .48, "E7": .32, "E8": .20, "B14": 193750},
                "Zahlungen": {"D11": 73574, "D15": 12000, "D16": 107100},
            },
            "gesellschafterstreit-zink-und-zunder": {
                "Bank und OP": {"D11": 19800, "D23": 26800, "D24": 51800},
                "Darlehen und Wünsche": {"D6": 60000, "D7": 10000, "D8": 70000},
            },
        }
        for case in CASES:
            path = ROOT / "testakten" / case["slug"] / ADDENDUM / case["xlsx"]
            workbook = load_workbook(path, data_only=True)
            try:
                for title, cells in expected[case["slug"]].items():
                    for coordinate, amount in cells.items():
                        self.assertAlmostEqual(workbook[title][coordinate].value, amount, places=8,
                                               msg=f"{case['slug']}/{title}/{coordinate}")
            finally:
                workbook.close()

    def test_skills_prompts_versions_and_links(self):
        skills = list((PLUGIN / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 11)
        for path in skills:
            text = path.read_text(encoding="utf-8")
            self.assertTrue(text.startswith("---\n"), str(path))
            frontmatter = yaml.safe_load(text.split("---", 2)[1])
            self.assertEqual(set(frontmatter), {"name", "description"})
            self.assertEqual(frontmatter["name"], path.parent.name)
            self.assertLessEqual(len(frontmatter["description"]), 360)
            self.assertTrue(path.parent.name.isascii())
            self.assertLessEqual(len(path.parent.name), 64)
            self.assertEqual(re.findall(r"^## ([1-6])\. ", text, re.M), list("123456"))
            output = text.split("## 5.", 1)[1].split("## 6.", 1)[0]
            for required in ("Times New Roman", "11 pt"):
                self.assertIn(required, output, str(path))
            self.assertRegex(output.lower(), r"ausformuliert|ausformulierungspflicht")
        self.assertEqual(validate_files(PLUGIN, PLUGIN.name, ROOT), [])
        for kind in ("werkstatt", "schnellstart", "hauptproblem"):
            md, txt = [PLUGIN / f"gesellschafterstreit-{kind}.{ext}" for ext in ("md", "txt")]
            self.assertEqual(md.read_bytes(), txt.read_bytes())
            if kind != "werkstatt":
                self.assertLessEqual(md.stat().st_size, 7500)
        for kind in (".claude-plugin", ".codex-plugin"):
            self.assertEqual(json.loads((PLUGIN / kind / "plugin.json").read_text())["version"], VERSION)
        market = json.loads((ROOT / ".claude-plugin/marketplace.json").read_text())
        self.assertEqual([entry["version"] for entry in market["plugins"] if entry["name"] == PLUGIN.name], [VERSION])
        routes = json.loads((ROOT / "scripts/scoped-release-assets.json").read_text())["assets"]
        self.assertEqual(routes["gesellschafterstreit.zip"], f"gesellschafterstreit-v{VERSION}")
        files = list(PLUGIN.rglob("*.md")) + [ROOT / "testakten" / case["slug"] / "README.md" for case in CASES]
        for path in files:
            for target in re.findall(r"\[[^\]]*\]\(([^)]+)\)", path.read_text(encoding="utf-8")):
                parsed = urlsplit(target.strip("<>"))
                if not parsed.scheme and not parsed.netloc and parsed.path:
                    self.assertTrue((path.parent / unquote(parsed.path)).resolve().exists(), f"{path}: {target}")

    def test_register_orders_originals_before_addenda_and_identifies_paths(self):
        builder = module("build-gesellschafter-vorfuehrpakete.py")
        for case in CASES:
            directory = ROOT / "testakten" / case["slug"]
            paths = builder.sources(directory)
            original_count = len(BASELINE_SHA256[case["slug"]])
            self.assertTrue(all(path.parent == directory for path in paths[:original_count]))
            self.assertTrue(all(path.parent.name == ADDENDUM for path in paths[original_count:]))
            combined = directory / "gesamt-pdf" / f"{case['slug']}_gesamt.pdf"
            data = combined.read_bytes()
            self.assertEqual(pdf_content_errors(data), [])
            reader = PdfReader(io.BytesIO(data))
            outlines = [item for item in reader.outline if isinstance(item, dict)]
            self.assertEqual([str(item["/Title"]) for item in outlines[1:]], [path.relative_to(directory).as_posix() for path in paths])
            positions = [reader.get_destination_page_number(item) for item in outlines[1:]]
            self.assertEqual(positions, sorted(set(positions)))
            register = "\n".join(page.extract_text() or "" for page in reader.pages[:positions[0]])
            for path in paths:
                self.assertIn(normalized(path.relative_to(directory).as_posix()), normalized(register))
            self.assertIn("02.10.2026", register)
            self.assertIn("08.10.2026", register)

    def test_release_assets_native_bytes_pdf_content_and_checksums(self):
        if DIST is None:
            self.skipTest("Paketprüfung nur mit --dist VERZEICHNIS")
        expected = {"gesellschafterstreit.zip", "checksums-sha256.txt", *(path.name for path in release_copies())}
        expected.update(f"testakte-{case['slug']}{suffix}.zip" for case in CASES for suffix in ("", "-einzelpdfs"))
        self.assertEqual(len(expected), 19)
        self.assertEqual({path.name for path in DIST.iterdir()}, expected)
        for path in release_copies():
            self.assertEqual((DIST / path.name).read_bytes(), path.read_bytes(), path.name)
        originals = module("validate-testakten-release-zips.py")
        singles = module("validate-testakten-einzelpdf-zips.py")
        for case in CASES:
            directory = ROOT / "testakten" / case["slug"]
            original_zip = DIST / f"testakte-{case['slug']}.zip"
            single_zip = DIST / f"testakte-{case['slug']}-einzelpdfs.zip"
            originals.assert_same(original_zip.name, originals.expected_entries(directory), originals.zip_entries(original_zip, require_notice=True))
            singles.assert_same(single_zip.name, ["README.txt", *singles.expected_arcnames(directory)], singles.zip_entries(single_zip, expected_suffix=".pdf"))
            for path in (original_zip, single_zip):
                with zipfile.ZipFile(path) as archive:
                    self.assertEqual(archive.namelist()[0], "README.txt")
                    self.assertEqual(archive.read("README.txt"), NOTICE_BYTES)
            with zipfile.ZipFile(original_zip) as archive:
                for source, name in working_dump_archive_pairs(directory, include_gesamt_pdf=True):
                    self.assertEqual(archive.read(name), source.read_bytes(), name)
            with zipfile.ZipFile(single_zip) as archive:
                for source, name in document_arcname_pairs(directory):
                    data = archive.read(name)
                    text = pdf_text(data)
                    if source.suffix == ".pdf":
                        self.assertEqual(data, source.read_bytes(), name)
                    elif source.suffix == ".docx":
                        body_text = docx_pdf_body_text(data)
                        for passage in docx_passages(source):
                            self.assert_passage(passage, body_text, f"Word-Inhalt fehlt: {name}")
                        self.assert_pdf_layout(data, name)
                    elif source.suffix == ".xlsx":
                        views = workbook_pdf_views(data)
                        workbook = load_workbook(source, data_only=True)
                        try:
                            for sheet in workbook:
                                for row in sheet:
                                    for cell in row:
                                        if cell.value is None or not str(cell.value).strip():
                                            continue
                                        candidates = numeric_texts(cell)
                                        self.assertTrue(any(normalized(value) in view for value in candidates for view in views),
                                                        f"Excel-Zellinhalt fehlt: {name}/{sheet.title}/{cell.coordinate}: {cell.value!r}")
                        finally:
                            workbook.close()
                        self.assert_pdf_layout(data, name)
                    elif source.suffix == ".eml":
                        message = BytesParser(policy=policy.default).parsebytes(source.read_bytes())
                        for passage in re.split(r"\n\s*\n", message.get_body(preferencelist=("plain",)).get_content()):
                            self.assert_passage(passage, text, f"Mailinhalt fehlt: {name}")
                    elif source.suffix == ".txt":
                        self.assert_passage(source.read_text(encoding="utf-8"), text, f"Textinhalt fehlt: {name}")
        with zipfile.ZipFile(DIST / "gesellschafterstreit.zip") as archive:
            expected_plugin = {path.relative_to(PLUGIN).as_posix() for path in PLUGIN.rglob("*")
                               if path.is_file() and not path.is_symlink() and "__pycache__" not in path.parts
                               and path.suffix != ".pyc" and path.name not in {".DS_Store", "CLAUDE.md"}
                               and not path.name.endswith(PROMPT_SUFFIXES)}
            self.assertCountEqual(archive.namelist(), expected_plugin)
            self.assertIsNone(archive.testzip())
            self.assertEqual(sum(name.endswith("/SKILL.md") for name in archive.namelist()), 11)
            for name in archive.namelist():
                self.assertEqual(archive.read(name), (PLUGIN / name).read_bytes(), name)
        names = []
        for line in (DIST / "checksums-sha256.txt").read_text(encoding="utf-8").splitlines():
            self.assertRegex(line, r"^[0-9a-f]{64}  [^/\\]+$")
            digest, name = line.split("  ", 1)
            names.append(name)
            self.assertEqual(hashlib.sha256((DIST / name).read_bytes()).hexdigest(), digest, name)
        self.assertCountEqual(names, expected - {"checksums-sha256.txt"})


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dist", type=Path)
    args, remaining = parser.parse_known_args()
    DIST = args.dist.resolve() if args.dist else None
    unittest.main(argv=[__file__, *remaining])
