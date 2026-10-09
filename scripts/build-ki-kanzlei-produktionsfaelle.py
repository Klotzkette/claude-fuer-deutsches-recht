#!/usr/bin/env python3
"""Setzt die drei einzeln redigierten Produktionsakten in native Originalformate."""
import csv
import importlib.util
import json
import tempfile
import zipfile
from datetime import datetime, timezone
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from testakte_office_pdf import render_office
from testakte_disclaimer import NOTICE_MARKDOWN

ROOT = Path(__file__).resolve().parents[1]
DATA = Path(__file__).with_name('ki-kanzlei-produktionsfaelle.json')


def load_module(name, file):
    spec = importlib.util.spec_from_file_location(name, ROOT / 'scripts' / file)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def word(path, item, case):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(2.3), Cm(2.3)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == 'Normal' else 14)
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after, normal.line_spacing = Pt(7), 1.12
    normal.widow_control = True
    normal.keep_together = True
    header = section.header.paragraphs[0]
    header.text = item['sender'].split('\n')[0] + ' | ' + case['matter']
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.text = case['matter'] + ' | Seite '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.add_paragraph(item['sender'])
    doc.add_paragraph(item['recipient'])
    doc.add_paragraph(item['date'])
    doc.add_paragraph(item['title'], 'Title')
    for paragraph in item['body']:
        lines = paragraph.split('\n', 1)
        if len(lines) == 2 and lines[0][:1].isdigit():
            doc.add_paragraph(lines[0], 'Heading 2')
            doc.add_paragraph(lines[1])
        else:
            doc.add_paragraph(paragraph)
    props = doc.core_properties
    props.author = props.last_modified_by = 'Klotzkette'
    props.title = item['title']
    props.created = props.modified = datetime(2026, 10, 9, tzinfo=timezone.utc)
    doc.save(path)
    # Stabile Containerdaten halten Wiederholungen auf den eigentlichen Inhalt begrenzt.
    with zipfile.ZipFile(path) as source:
        parts = [(name, source.read(name)) for name in source.namelist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as target:
        for name, data in parts:
            info = zipfile.ZipInfo(name, (2026, 10, 9, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)


def main():
    section = load_module('ki_download_section', 'inject-gesamt-pdf-section.py')
    for case in json.loads(DATA.read_text()):
        folder = ROOT / 'testakten' / case['slug']
        folder.mkdir(parents=True, exist_ok=True)
        for item in case['documents']:
            path = folder / item['file']
            if path.suffix == '.docx':
                word(path, item, case)
            elif path.suffix == '.pdf':
                with tempfile.TemporaryDirectory(prefix='kanzlei-beleg-') as temp:
                    source = Path(temp) / (path.stem + '.docx')
                    word(source, item, case)
                    data = render_office(source)
                    if data is None:
                        raise RuntimeError('Native Office-Konvertierung erforderlich: ' + path.name)
                    path.write_bytes(data)
            elif path.suffix == '.eml':
                msg = EmailMessage(policy=SMTP)
                msg['From'], msg['To'] = item['from'], item['to']
                msg['Date'] = format_datetime(datetime.fromisoformat(item['date']))
                msg['Subject'] = item['subject']
                msg['Message-ID'] = f"<{case['slug']}.{path.stem}@post.example>"
                msg.set_content(item['body'], charset='utf-8')
                path.write_bytes(msg.as_bytes())
            elif path.suffix == '.csv':
                with path.open('w', encoding='utf-8', newline='') as handle:
                    csv.writer(handle, delimiter=';').writerows(item['rows'])
            elif path.suffix == '.txt':
                path.write_text(item['body'], encoding='utf-8')
            else:
                raise ValueError(path)
        overview = '\n'.join(f"| [{d['file']}]({d['file']}) | {d.get('title', d.get('subject', Path(d['file']).stem.replace('_', ' ')))} |" for d in case['documents'])
        text = (f"# {case['title']}\n\n{case['summary']}\n\n"
                f"Plugin: [KI-native Kanzlei](../../ki-native-kanzlei/README.md). Aktenstand: 9. Oktober 2026. Interne Kennung: `{case['matter']}`.\n\n"
                "## 1. Originalunterlagen\n\nZehn einzeln lesbare Unterlagen. Behauptungen, Erinnerungen und nicht beigefügte Dokumente bleiben als solche erkennbar. Die Akte enthält keine Musterlösung.\n\n"
                + NOTICE_MARKDOWN + "\n\n| Datei | Inhalt |\n| --- | --- |\n" + overview + '\n\n' +
                section.section_block(case['slug'], f"gesamt-pdf/{case['slug']}_gesamt.pdf", has_einzelpdf=True) +
                "\n\n## 4. Herkunft und Verwendung\n\nPersonen, Unternehmen, Aktenzeichen und Vorgänge sind erfunden. Behördenbezeichnungen dienen der Einordnung; die Unterlagen sind keine tatsächlich ergangenen gerichtlichen Dokumente. Die reservierten Kontaktadressen sind nicht für echten Versand bestimmt. Originale und beide PDF-Ausgaben bilden denselben Aktenstand ab. Die Archive sind flach und enthalten den zweisprachigen Hinweis in README.txt.\n\n<!-- reserved-example-contacts -->\n\n"
                "[Alle Testakten](../README.md) · [Startseite](../../README.md) · [Downloads](../../ASSET_INDEX.md)\n")
        (folder / 'README.md').write_text(text, encoding='utf-8')
        rubric = {'name': case['title'], 'plugin': 'ki-native-kanzlei', 'stand': '2026-10-09', 'checks': [
            {'id': 'bestand', 'check_type': 'working_file_count', 'description': 'Zehn native Originalunterlagen.', 'min': 10},
            *[{'id': f'original-{i}', 'check_type': 'file_exists', 'description': d['file'], 'path': d['file']} for i, d in enumerate(case['documents'], 1)],
            {'id': 'mandatsarbeit', 'check_type': 'human_review', 'description': 'Belege, Behauptungen, fehlende Unterlagen und konkrete Freigaben werden auseinandergehalten; ein ausformuliertes Arbeitsprodukt entsteht.'}]}
        (folder / 'rubric.yaml').write_text(json.dumps(rubric, ensure_ascii=False, indent=2) + '\n')
        print(case['slug'], len(case['documents']), flush=True)


if __name__ == '__main__':
    main()
