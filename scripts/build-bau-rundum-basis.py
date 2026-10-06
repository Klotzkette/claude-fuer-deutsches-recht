#!/usr/bin/env python3
"""Rendert ausschließlich die fünf individuell verfassten Bau-Einstiegsakten.

Aufruf mit dem gebündelten Python und PYTHONPATH laut Aufgabenübergabe.
--render-check exportiert DOCX nur mit der gebündelten Büro-Runtime und erzeugt
temporäre Seitenbilder innerhalb des jeweiligen Aktenordners für die Sichtprüfung.
--new-only ergänzt Originale und Metadaten, ohne bestehende Originale zu verändern.
--clean-qa entfernt ausschließlich diese eigenen Prüfdateien.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import mimetypes
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from email.utils import format_datetime, getaddresses
from pathlib import Path
from xml.sax.saxutils import escape

sys.dont_write_bytecode = True
from bau_rundum_basis_daten import AKTEN
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, PngImagePlugin
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


REPO = Path(__file__).resolve().parent.parent
SOFFICE = Path('/Users/klotzkette/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
FONT_DIR = Path('/System/Library/Fonts/Supplemental')
NOTICE = ("> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n>\n"
          "> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.")
STYLES = {}


def setup_fonts():
    for name, filename in [('TNR', 'Times New Roman.ttf'), ('TNRB', 'Times New Roman Bold.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(FONT_DIR / filename)))
    STYLES.update(
        body=ParagraphStyle('body', fontName='TNR', fontSize=11, leading=14, spaceAfter=7),
        title=ParagraphStyle('title', fontName='TNRB', fontSize=17, leading=20, spaceAfter=12),
        meta=ParagraphStyle('meta', fontName='TNR', fontSize=9, leading=11, spaceAfter=5),
        head=ParagraphStyle('head', fontName='TNRB', fontSize=12, leading=15, spaceBefore=9, spaceAfter=9, keepWithNext=True),
        cell=ParagraphStyle('cell', fontName='TNR', fontSize=10, leading=12, alignment=TA_LEFT),
        cellhead=ParagraphStyle('cellhead', fontName='TNRB', fontSize=10, leading=12),
    )


def paragraphs(page):
    for block in page if isinstance(page, list) else [page]:
        if isinstance(block, dict):
            yield block
        else:
            yield from (s.strip() for s in block.strip().split('\n\n') if s.strip())


def heading(text):
    return bool(re.match(r'^\d+(?:\.\d+)*\s+[^\n]+$', text))


def pdf_table(block):
    rows = [[Paragraph(escape(cell).replace('\n', '<br/>'), STYLES['cellhead' if i == 0 else 'cell'])
             for cell in row] for i, row in enumerate([block['kopf']] + block['zeilen'])]
    table = Table(rows, colWidths=[v * 168 * mm for v in block['breiten']], repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e6edf1')),
        ('GRID', (0, 0), (-1, -1), 0.45, colors.HexColor('#d9d9d9')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    return table


def write_pdf(path, item):
    story = []
    for n, page in enumerate(item['seiten']):
        if n:
            story.append(PageBreak())
        else:
            story.append(Paragraph(escape(item['titel']), STYLES['title']))
            for line in [item['absender'], 'An: ' + item['empfaenger'], item['datum'] + ' | ' + item['zeichen']]:
                story.append(Paragraph(escape(line), STYLES['meta']))
            story.append(Spacer(1, 7 * mm))
        for block in paragraphs(page):
            if isinstance(block, dict):
                story.extend([pdf_table(block), Spacer(1, 4 * mm)])
            else:
                story.append(Paragraph(escape(block).replace('\n', '<br/>'), STYLES['head' if heading(block) else 'body']))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('TNR', 9)
        canvas.drawString(21 * mm, 13 * mm, item['zeichen'].split(' / ')[0])
        canvas.drawRightString(189 * mm, 13 * mm, f'Seite {doc.page}')
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=(210 * mm, 297 * mm), leftMargin=21 * mm,
                            rightMargin=21 * mm, topMargin=18 * mm, bottomMargin=22 * mm,
                            title=item['titel'], author=item['absender'], subject=item['zeichen'])
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    actual = len(PdfReader(path).pages)
    if actual != len(item['seiten']):
        raise RuntimeError(f'{path.name}: erwartet {len(item["seiten"])} Seiten, erzeugt {actual}')


def word_table(doc, block):
    table = doc.add_table(rows=1, cols=len(block['kopf']))
    table.autofit = False
    for i, width in enumerate(block['breiten']):
        table.columns[i].width = Cm(16.8 * width)
    for i, row in enumerate([block['kopf']] + block['zeilen']):
        cells = table.rows[0].cells if i == 0 else table.add_row().cells
        for j, value in enumerate(row):
            cells[j].width = Cm(16.8 * block['breiten'][j])
            cells[j].text = value
            pr = cells[j]._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                element = OxmlElement(f'w:{edge}')
                for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    element.set(qn(f'w:{k}'), v)
                borders.append(element)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement(f'w:{edge}')
                e.set(qn('w:w'), '85')
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            pr.append(margins)
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = i == 0
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E6EDF1')
                pr.append(shade)
        trpr = table.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def write_docx(path, item):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.1)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Footer']:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Times New Roman', Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.08
    doc.styles['Title'].font.size = Pt(17)
    doc.styles['Title'].font.bold = True
    for name in ['Heading 1', 'Heading 2']:
        doc.styles[name].font.size = Pt(12)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_before = Pt(9)
        doc.styles[name].paragraph_format.space_after = Pt(9)
    doc.core_properties.author = item['absender']
    doc.core_properties.last_modified_by = item['absender'].split(' | ')[0]
    doc.core_properties.title = item['titel']
    doc.core_properties.subject = item['zeichen']
    doc.core_properties.language = 'de-DE'
    doc.core_properties.comments = ''
    date = datetime.strptime(item['datum'][:10], '%d.%m.%Y').replace(tzinfo=timezone.utc)
    doc.core_properties.created = doc.core_properties.modified = date
    for n, page in enumerate(item['seiten']):
        if n:
            doc.add_page_break()
        else:
            doc.add_paragraph(item['titel'], 'Title')
            for line in [item['absender'], 'An: ' + item['empfaenger'], item['datum'] + ' | ' + item['zeichen']]:
                p = doc.add_paragraph(line)
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(9)
        for block in paragraphs(page):
            if isinstance(block, dict):
                word_table(doc, block)
            else:
                doc.add_paragraph(block, 'Heading 1' if heading(block) else 'Normal')
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(item['zeichen'].split(' / ')[0] + ' | Seite ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    p._p.append(field)
    for run in p.runs:
        run.font.size = Pt(9)
    # Remove inherited template rules and theme fonts, which override explicit names.
    for root in [doc.styles.element, doc.element]:
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        for fonts in root.iter(qn('w:rFonts')):
            for key in list(fonts.attrib):
                if key.endswith('Theme'):
                    del fonts.attrib[key]
    doc.save(path)


def draw_image(path, item):
    im = Image.new('RGB', (1500, 1100), 'white')
    d = ImageDraw.Draw(im)
    font = ImageFont.truetype(str(FONT_DIR / 'Arial.ttf'), 27)
    small = ImageFont.truetype(str(FONT_DIR / 'Arial.ttf'), 23)
    title = ImageFont.truetype(str(FONT_DIR / 'Arial Bold.ttf'), 36)
    d.text((55, 35), item['titel'], font=title, fill='#17252e')
    d.text((55, 95), item['datum'], font=font, fill='#354750')
    d.text((55, 145), 'Illustration / Skizze, keine Baustellenfotografie. Nicht maßstäblich.', font=font, fill='#9b3e20')
    blue, red = '#316679', '#b03b35'
    kind = item['art']
    if kind == 'fenster':
        d.rectangle((250, 250, 1040, 790), fill='#e8eff1', outline=blue, width=8)
        d.rectangle((290, 290, 1000, 740), fill='#f8fbfc', outline=blue, width=7)
        d.line((645, 290, 645, 740), fill=blue, width=7)
        d.rectangle((250, 780, 1040, 815), fill='#d7d9d9', outline=blue, width=3)
        d.line((825, 798, 1010, 798), fill=red, width=15)
        d.line((930, 802, 1110, 670), fill=red, width=4)
        d.text((1075, 585), 'Band / Putz', font=font, fill=red)
        d.text((1075, 625), 'offen', font=font, fill=red)
        d.text((315, 410), 'F07', font=title, fill=blue)
        d.text((330, 845), 'Raumseite / unterer Anschluss', font=font, fill=blue)
    elif kind == 'ablauf':
        d.rectangle((250, 250, 1180, 850), fill='#f1f3f3', outline=blue, width=6)
        for x in range(405, 1180, 155):
            d.line((x, 250, x, 850), fill='#c9d0d2', width=2)
        for y in range(400, 850, 150):
            d.line((250, y, 1180, y), fill='#c9d0d2', width=2)
        d.ellipse((580, 490, 800, 670), fill='#a9dce9', outline=blue, width=3)
        d.rectangle((790, 530, 890, 630), fill='#fafafa', outline='#465057', width=5)
        for x in range(800, 889, 15):
            d.line((x, 535, x, 625), fill='#465057', width=3)
        d.text((870, 690), 'Rost', font=font, fill=blue)
        d.text((440, 690), 'Wasserrest', font=font, fill=blue)
        d.text((280, 275), 'Nordwand', font=font, fill=blue)
        d.text((265, 805), 'West', font=font, fill=blue)
    elif kind == 'leuchten':
        d.rectangle((400, 230, 1060, 870), outline=blue, width=6)
        d.rectangle((413, 244, 1047, 283), fill='#d6dfd6')
        d.text((450, 285), 'Nord: feste Regale', font=small, fill=blue)
        for col in range(6):
            d.text((437 + col * 103, 205), str(col + 1), font=small, fill=blue)
            for row in range(6):
                x, y = 435 + col * 103, 335 + row * 78
                d.rectangle((x, y, x + 55, y + 38), fill='#ffedb0', outline='#8b7030', width=3)
        d.text((65, 445), 'West', font=font, fill=blue)
        d.text((1150, 445), 'Ost', font=font, fill=blue)
        d.text((440, 820), 'Südeingang / Bedienstelle', font=small, fill=blue)
    elif kind == 'raeume':
        for box, label in [((200, 280, 695, 650), '0.11\n43,20 m²'), ((695, 280, 1150, 650), '0.12\n28,80 m²'), ((200, 650, 1150, 800), '0.10 Flur: 24,00 m²'), ((1150, 650, 1420, 800), '0.13\n12,00 m²')]:
            d.rectangle(box, fill='#edf2f3', outline=blue, width=5)
            d.multiline_text((box[0] + 25, box[1] + 28), label, font=font, fill=blue, spacing=12)
        d.line((695, 280, 695, 650), fill=red, width=10)
        d.text((205, 220), 'Raumzuordnung, keine maßstäbliche Wandgeometrie', font=font, fill=blue)
        d.text((880, 665), 'T12', font=small, fill=red)
    elif kind == 'felder':
        d.rectangle((180, 280, 1290, 790), outline=blue, width=6)
        d.rectangle((190, 290, 730, 780), fill='#e5eee2')
        d.rectangle((745, 290, 1280, 780), fill='#fff0e8')
        d.line((740, 280, 740, 790), fill=blue, width=5)
        d.text((250, 350), 'Westfeld\nSchalung vorbereitet', font=font, fill=blue, spacing=15)
        d.text((840, 350), 'Ostfeld\nBewehrung abschnittsweise', font=font, fill=blue, spacing=15)
        d.rectangle((1070, 475, 1160, 740), fill='#d9c5b5', outline=red, width=4)
        d.text((850, 805), 'Technikrinne bei Achsen 4 bis 5', font=small, fill=red)
        for x, label in [(180, '1'), (740, '3'), (1020, '4'), (1290, '5')]:
            d.text((x - 10, 235), label, font=font, fill=blue)
    elif kind == 'bohrungen':
        d.rectangle((230, 260, 1210, 845), outline=blue, width=6)
        d.rectangle((1010, 270, 1200, 835), fill='#fff0d5')
        for x, y, label in [(380, 400, 'KB1'), (660, 680, 'KB2'), (810, 400, 'KB3'), (1100, 610, 'KB4')]:
            d.ellipse((x - 13, y - 13, x + 13, y + 13), fill=red)
            d.text((x - 30, y + 28), label, font=font, fill=blue)
        d.text((260, 285), 'West', font=font, fill=blue)
        d.text((1040, 285), 'Ost', font=font, fill=blue)
        d.text((490, 860), '24,00 m Hallenlänge', font=font, fill=blue)
        d.text((1020, 745), 'erst 14.09.\nzugänglich', font=small, fill=blue, spacing=10)
    else:
        raise ValueError(kind)
    d.text((55, 935), item['beschriftung'], font=font, fill='#17252e')
    d.text((55, 992), item['untertitel'], font=small, fill='#354750')
    info = PngImagePlugin.PngInfo()
    info.add_text('Title', item['titel'])
    info.add_text('Description', 'Illustration / Skizze, keine Baustellenfotografie. ' + item['untertitel'])
    info.add_text('Source', item['datum'])
    im.save(path, pnginfo=info, optimize=True)


def message_id(slug, filename):
    return f'<{hashlib.sha256((slug + filename).encode()).hexdigest()[:24]}@aktenpost.example>'


def write_email(folder, slug, item, attachments):
    message = EmailMessage(policy=policy.SMTP)
    for key, value in [('From', item['absender']), ('To', item['empfaenger']),
                       ('Date', format_datetime(datetime.fromisoformat(item['datum']))),
                       ('Subject', item['betreff']), ('Message-ID', message_id(slug, item['datei']))]:
        message[key] = value
    if item.get('antwort'):
        message['In-Reply-To'] = message_id(slug, item['antwort'])
    body = item['text'].strip()
    if attachments:
        body += '\n\nAnlagen:\n' + '\n'.join(attachments)
    message.set_content(body, charset='utf-8')
    for name in attachments:
        content_type = mimetypes.guess_type(name)[0] or 'application/octet-stream'
        major, minor = content_type.split('/', 1)
        message.add_attachment((folder / name).read_bytes(), maintype=major, subtype=minor, filename=name)
    (folder / item['datei']).write_bytes(message.as_bytes())


def write_meta(folder, slug, case, originals):
    intro = f'''<!-- decimal-headings -->
# 1. {case['titel']}

Autor: Klotzkette. Aktenstand: {case['stand']}. Zugeordnetes Plugin: `bauwirtschaft-anfaenger`.

{case['beschreibung']}

## 1.1. Herkunft und Abgrenzung

Alle Personen, Unternehmen, Stellen, Anschriften, Projekte und Zahlen dieser Akte sind neu erfunden. Der Ortsname ist real. Kontaktadressen verwenden ausschließlich reservierte `.example`-Domains. Bilddateien sind beschriftete Illustrationen beziehungsweise technische Skizzen, keine Baustellenfotografien. Die Unterlagen enthalten keine Musterlösung. `rubric.yaml` bleibt außerhalb der Archive. Kein Live-Modelltest wurde durchgeführt.
<!-- reserved-example-contacts -->

## 1.2. Akte herunterladen

Das Gesamt-PDF dient zum Lesen und Ausdrucken. Das Originalformat-ZIP enthält die bearbeitbaren Unterlagen und E-Mails; das Einzel-PDF-ZIP enthält jede Unterlage als separate PDF-Datei.

{NOTICE}

| Fassung | Datei |
| --- | --- |
| Gesamt-PDF | [Gesamt-PDF](gesamt-pdf/{slug}_gesamt.pdf) |
| Originalformat-ZIP | [Originaldateien](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/bauwirtschaft-rundum-v445.33.2/testakte-{slug}.zip) |
| Einzel-PDF-ZIP | [Einzel-PDFs](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/bauwirtschaft-rundum-v445.33.2/testakte-{slug}-einzelpdfs.zip) |

## 1.3. Originalunterlagen

{len(originals)} eigenständige Originaldateien; E-Mail-Anhänge sind bytegleiche Kopien bereits aufgeführter Originale und werden nicht zusätzlich gezählt.

{NOTICE}

| Datei | Format |
| --- | --- |
'''
    intro += ''.join(f'| [{name}]({name}) | {Path(name).suffix[1:].upper()} |\n' for name in sorted(originals))
    intro += '''
## 1.4. Erzeugung und Qualitätssicherung

Die Texte stehen individuell verfasst in `scripts/bau_rundum_basis_daten.py`. `scripts/build-bau-rundum-basis.py` rendert ausschließlich diese fünf Originalakten. Technische und rechtliche Aussagen sind innerhalb der jeweiligen Projektunterlagen zu prüfen; die Akte enthält keine allgemeine Normsammlung oder technische Bemessungsfreigabe. Die Rubrik dient nur der internen Auswertung.
'''
    (folder / 'README.md').write_text(intro, encoding='utf-8')
    # JSON scalar syntax is valid YAML and preserves German text without quoting ambiguities.
    import json
    scalar = lambda value: json.dumps(value, ensure_ascii=False)
    yaml = f'name: {slug}\nplugin: bauwirtschaft-anfaenger\nauthor: Klotzkette\ndescription: {scalar(case["beschreibung"])}\nchecks:\n'
    for n, text in enumerate(case['pruefung'], 1):
        yaml += f'  - id: fachpruefung-{n}\n    check_type: human_review\n    description: {scalar(text)}\n'
    for n, name in enumerate(sorted(originals), 1):
        yaml += f'  - id: original-{n:02d}\n    check_type: file_exists\n    path: {scalar(name)}\n    description: {scalar("Eigenständige Originalunterlage " + name)}\n'
    (folder / 'rubric.yaml').write_text(yaml, encoding='utf-8')


def validate(folder, originals):
    assert len(originals) >= 10
    assert len({Path(n).suffix for n in originals}) >= 4
    for name in originals:
        path = folder / name
        if path.suffix == '.eml':
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            for key in ['From', 'To', 'Date', 'Subject', 'Message-ID', 'MIME-Version']:
                assert message[key], (name, key)
            assert message.get_body(preferencelist=('plain',)), name
            text = message.get_body(preferencelist=('plain',)).get_content()
            for part in message.iter_attachments():
                assert part.get_payload(decode=True) == (folder / part.get_filename()).read_bytes(), name
            for _, address in getaddresses([message['From'], message['To']]):
                assert address.endswith('.example'), address
        elif path.suffix == '.pdf':
            text = '\n'.join(p.extract_text() for p in PdfReader(path).pages)
        elif path.suffix == '.docx':
            doc = Document(path)
            text = '\n'.join(p.text for p in doc.paragraphs)
            text += '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        elif path.suffix == '.csv':
            with path.open(encoding='utf-8-sig', newline='') as stream:
                rows = list(csv.reader(stream, delimiter=';'))
            assert len(rows) >= 5, name
            assert all(len(r) == len(rows[0]) for r in rows), name
            continue
        else:
            continue
        assert len(text.strip()) >= 600, (name, len(text))
        assert 'Diese Testakte wurde mit KI' not in text, name
        assert '\ufffd' not in text, name


def build(slug, new_only=False):
    case = AKTEN[slug]
    folder = REPO / 'testakten' / slug
    folder.mkdir(parents=True, exist_ok=True)
    prior_hashes = {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in folder.iterdir()
        if new_only and p.is_file() and p.suffix in {'.docx', '.pdf', '.eml', '.csv', '.txt', '.png'}
    }

    def should_write(path):
        return not new_only or not path.exists()

    originals = []
    for item in case['dokumente']:
        path = folder / item['datei']
        if should_write(path):
            (write_docx if path.suffix == '.docx' else write_pdf)(path, item)
        originals.append(path.name)
    for key in ['texte', 'zusatztexte']:
        for name, text in case.get(key, {}).items():
            if should_write(folder / name):
                (folder / name).write_text(text, encoding='utf-8')
            originals.append(name)
    for name, (header, rows) in case.get('csv', {}).items():
        if should_write(folder / name):
            with (folder / name).open('w', encoding='utf-8', newline='') as stream:
                writer = csv.writer(stream, delimiter=';')
                writer.writerows([header] + rows)
        originals.append(name)
    for item in case['bilder']:
        if should_write(folder / item['datei']):
            draw_image(folder / item['datei'], item)
        originals.append(item['datei'])
    attachments = {
        'bau-rundum-begehung-einbeck': ['03_Randnotizen_Lenz.txt', '05_Gewerkekarte.pdf'],
        'bau-rundum-lv-abgleich-detmold': ['03_Baubeschreibung_02.docx', '04_LV_Ausbau_03.docx'],
        'bau-rundum-bieterfragen-celle': ['03_Leistungsbeschreibung_02.docx', '05_Produktnotiz.docx'],
        'bau-rundum-behinderung-soest': ['02_Diktat_Kroll.txt', '04_Ablauf_01.docx', '11_Ablauf_02.docx'],
        'bau-rundum-baugrund-verden': ['02_Baugrundbericht.pdf', '03_Ergaenzung_01.pdf', '06_Laborwerte.csv', '07_Wasserbeobachtungen.csv'],
    }
    for i, item in enumerate(case['mails']):
        mail_attachments = item.get('anlagen')
        if mail_attachments is None:
            mail_attachments = attachments[slug] if i == 0 else []
        if should_write(folder / item['datei']):
            write_email(folder, slug, item, mail_attachments)
        originals.append(item['datei'])
    write_meta(folder, slug, case, originals)
    validate(folder, originals)
    for name, digest in prior_hashes.items():
        assert hashlib.sha256((folder / name).read_bytes()).hexdigest() == digest, (slug, name, 'Altoriginal verändert')
    print(f'{slug}: {len(originals)} Originale, {len({Path(n).suffix for n in originals})} Dateitypen; Struktur und Anhänge geprüft.', flush=True)
    if new_only:
        print(f'{slug}: {len(prior_hashes)} vorhandene Originale bytegleich erhalten.', flush=True)
    return folder


def render_check(folder, slug):
    import pypdfium2
    import pdfplumber
    qa = folder / '.basis-qa'
    qa.mkdir(exist_ok=True)
    docs = sorted(folder.glob('*.docx'))
    profile = qa / 'office-profile'
    command = [str(SOFFICE), '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(qa)]
    result = subprocess.run(command + [str(p) for p in docs], capture_output=True, text=True, timeout=180)
    if result.returncode:
        raise RuntimeError(result.stdout + result.stderr)
    expected = {Path(d['datei']).stem: len(d['seiten']) for d in AKTEN[slug]['dokumente']}
    pdfs = sorted(folder.glob('*.pdf')) + sorted(qa.glob('*.pdf'))
    rendered = []
    for path in pdfs:
        if path.parent == qa:
            assert path.stem in expected
        pdf = pypdfium2.PdfDocument(path)
        assert len(pdf) == expected[path.stem], (path.name, len(pdf), expected[path.stem])
        with pdfplumber.open(path) as checked:
            for page in checked.pages:
                for word in page.extract_words():
                    assert word['x0'] >= 25 and word['x1'] <= page.width - 25, (path.name, 'horizontaler Überlauf', word)
                    assert word['top'] >= 15 and word['bottom'] <= page.height - 15, (path.name, 'vertikaler Überlauf', word)
        for n, page in enumerate(pdf, 1):
            name = f'{path.stem}-Seite-{n}.png'
            page.render(scale=1.25).to_pil().convert('RGB').save(qa / name)
            page.close()
            rendered.append(name)
        pdf.close()
    for index in range(0, len(rendered), 4):
        sheet = Image.new('RGB', (1500, 2200), '#dce2e5')
        draw = ImageDraw.Draw(sheet)
        font = ImageFont.truetype(str(FONT_DIR / 'Arial.ttf'), 18)
        for j, name in enumerate(rendered[index:index + 4]):
            with Image.open(qa / name) as source:
                source.thumbnail((725, 1035))
                x, y = (j % 2) * 750, (j // 2) * 1100
                sheet.paste(source, (x + 12, y + 40))
                draw.text((x + 12, y + 12), name, font=font, fill='black')
        sheet.save(qa / f'Sichtbogen-{index // 4 + 1:02d}.png', optimize=True)
    if profile.exists():
        shutil.rmtree(profile)
    print(f'{slug}: {len(pdfs)} PDF/DOCX-Originale, {len(rendered)} Seiten gerendert; Seitenzahl und Textgrenzen geprüft.', flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('slugs', nargs='*', choices=None)
    parser.add_argument('--render-check', action='store_true')
    parser.add_argument('--new-only', action='store_true')
    parser.add_argument('--clean-qa', action='store_true')
    parser.add_argument('--check-only', action='store_true')
    args = parser.parse_args()
    slugs = args.slugs or list(AKTEN)
    if any(s not in AKTEN for s in slugs):
        parser.error('Nur die fünf zugewiesenen Akten sind zulässig.')
    setup_fonts()
    for slug in slugs:
        folder = REPO / 'testakten' / slug
        if args.clean_qa:
            qa = folder / '.basis-qa'
            if qa.exists():
                shutil.rmtree(qa)
            continue
        if not args.check_only:
            build(slug, new_only=args.new_only)
        if args.render_check:
            render_check(folder, slug)


if __name__ == '__main__':
    main()
