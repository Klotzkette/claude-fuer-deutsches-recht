#!/usr/bin/env python3
"""Erzeugt die Erfurter Forderungskaufakte aus einem deterministischen Fallmodell.

DOCX und PDF sind Originalunterlagen. Die Excel wird mit dem benachbarten
Artifact-Tool-Builder erzeugt; dieser Builder verwendet keine Excel-Ersatzbibliothek.
"""
from __future__ import annotations
import argparse
import calendar
import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

ROOT = Path(__file__).resolve().parents[1]
SLUG = "dd-bankfiliale-verbraucherdarlehen-erfurt"
OUT = ROOT / "testakten" / SLUG
MODEL = ROOT / "scripts/data/dd-bankakte.json"
BANK = "Ashcombe Personal Bank plc, Niederlassung Erfurt"
BUYER = "Ilmgrund Kreditservice GmbH"
ADDR = "Lindenbogen 18, 99084 Erfurt"
WARN = "Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\nThis test case file was generated with AI and is an experiment. Use at your own responsibility and risk."


def cent(v):
    return int((Decimal(str(v)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def money(v):
    s = f"{v / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{s} EUR"


# id, borrower, street, occupation, purpose, initial EUR, annual rate, topic,
# letter heading, letter sender role, letter full text, customer's distinct reply.
SEEDS = [
    (1, "Kunigunde Schlehenried", "Hollergrund 7, 99092 Erfurt", "Floristin", "Einbauküche und Haushaltsgeräte", 18000, 0.048, "Verbraucherinsolvenz",
     "Forderungsanmeldung und Auszug aus der Arbeitsfassung der Tabelle", "Dr. Malte Hohenacker, Insolvenzverwaltung",
     "Sehr geehrte Frau Voss, seit der Verfahrenseröffnung am 15.06.2026 führen wir die Korrespondenz für Frau Schlehenried. Ihre Anmeldung vom 26.06.2026 nennt die Darlehenshauptforderung und Nebenforderungen getrennt. Der anliegende Tabellenauszug gibt unsere Arbeitsfassung zum 24.09.2026 wieder; er ersetzt weder einen beglaubigten Tabellenauszug noch eine Vollstreckungsgrundlage. Die Schuldnerin bestreitet die Mahnkosten von 36 Euro. Ihre Forderung ist in unserer Liste als ungesichert vermerkt. Bitte teilen Sie uns vor einem Gläubigerwechsel die vollständige Abtretungskette und eine ladungsfähige Anschrift mit. Eine Quote können wir noch nicht zusagen. Die bei Ihnen nach dem Eröffnungsdatum weiterlaufenden Buchungszinsen sind nicht Bestandteil unserer vorliegenden Anmeldung. Bitte übersenden Sie eine Überleitung vom Anmeldebetrag zu Ihrem Septemberkontostand. Die Lohnabtretung wurde im Antrag lediglich erwähnt; eine unterschriebene Urkunde liegt hier nicht vor.",
     "Meine Schwester hat mir bei den Briefen geholfen. Seit Juni zahle ich nichts mehr direkt an die Bank. Bitte schicken Sie nicht wieder eine Mahnung an meinen früheren Laden. Die 36 Euro habe ich bestritten, weil ich die beiden Briefe nie erhalten habe. Ich bin nicht damit einverstanden, dass in einer Käuferliste bereits eine feste Insolvenzquote steht."),
    (2, "Balthasar Rummelspacher", "Amselbogen 21, 99089 Erfurt", "Lagerist", "Umschuldung zweier Ratenkäufe", 24000, 0.054, "Restschuldbefreiung und Altbestand",
     "Abgleich der Altvertragsforderung nach Abschluss des Insolvenzverfahrens", "Anke Zirbel, Schuldnerberatung",
     "Sehr geehrte Damen und Herren, Herr Rummelspacher hat uns Ihre Saldenbestätigung vom 18.09.2026 vorgelegt. Sie betrifft APB-0002, einen Vertrag vom September 2023. Das Insolvenzverfahren wurde am 12.02.2024 eröffnet; die auf den Altvertrag bezogene Restschuldbefreiung wurde nach dem uns übergebenen Beschluss vom 12.08.2026 im vorzeitig beendeten Verfahren erteilt. Die Akte enthält eine Erklärung sämtlicher angemeldeter Gläubiger zur Erledigung der Verfahrenskosten und Zustimmung zum vorzeitigen Abschluss. Der Beschluss ist nach dem Vermerk der Geschäftsstelle seit 04.09.2026 rechtskräftig. Ihre Bank hatte 2024 lediglich eine gewöhnliche Vertragsforderung angemeldet; ein besonderer Deliktsgrund ist in der Kopie nicht bezeichnet. Wir bitten um Erklärung, warum Ihre Septemberliste die Forderung ohne Sperrvermerk zur weiteren Einziehung vorsieht. Bitte unterscheiden Sie Ihre technische Forderungsbuchung von einem nach Ihrer Ansicht noch durchsetzbaren Anspruch. Bis zur Klärung erwarten wir, dass Sie keine erneute Zahlungsaufforderung veranlassen.",
     "Ich dachte, mit dem Schreiben vom Gericht sei das Thema erledigt. Ein Kollege von Ihnen sagte am Telefon, bei einem Verkauf fange die Forderung wieder neu an. Bitte erklären Sie das schriftlich. Ich kann den Gerichtsbeschluss am Montag noch einmal vollständig schicken; die Bank hat bisher nur das Deckblatt eingescannt."),
    (3, "Friedolin Kesselhut", "Wiesenbogen 9, 99094 Erfurt", "Busfahrer", "Gebrauchtwagen", 21000, 0.06, "Zahlungsrückstand und Kündigung",
     "Zahlungsaufforderung und angekündigte Gesamtfälligstellung", "Elin Voss, Ashcombe Forderungsservice",
     "Sehr geehrter Herr Kesselhut, nach unserer Buchung sind die Raten für Juli, August und September 2026 ausgeblieben. Wir hatten Sie mit unserem Schreiben vom 20.08.2026 zur Zahlung bis zum 07.09.2026 aufgefordert und für den Fall weiterer Nichtzahlung eine Beendigung des Darlehens angekündigt. Eine telefonische Besprechung der Einkommenssituation haben wir für den 25.08.2026 angeboten. Unter dem 22.09.2026 hat unser Team die Kündigung erklärt. Diese Nachricht ist nach unserer Versandliste als einfacher Brief aufgegeben worden; ein Einlieferungsbeleg fehlt. Der Ausdruck des Augustschreibens enthält in der archivierten Fassung keine aufaddierte Rückstandssumme, sondern nur die beiden damaligen Einzelraten. Bitte teilen Sie uns mit, ob Ihnen beide Schreiben zugegangen sind. Bis zum Abschluss der Prüfung führen wir die laufende Vertragsbuchung weiter. Der Kontosaldo im System ist deshalb nicht gleichbedeutend mit einer geprüften sofort fälligen Gesamtforderung. Ein Gespräch über eine tragfähige Ratenvereinbarung bleibt möglich.",
     "Die Schichten sind ausgefallen, deshalb konnte ich im Sommer nicht vollständig zahlen. Ein Augustschreiben habe ich bekommen, die Kündigung nicht. Am 29.09. habe ich 300 Euro überwiesen; das steht nicht auf Ihrem Ausdruck vom Vormittag. Bitte rufen Sie nicht bei meinem Arbeitgeber an. Ich möchte die Septemberrate und die neue Vereinbarung auseinanderhalten."),
    (4, "Meryem Aydin", "Ahornhof 11, 99086 Erfurt", "Physiotherapeutin", "Wohnungseinrichtung", 15000, 0.048, "Zustellung und Adresswechsel",
     "Rücklauf einer Zahlungsaufforderung", "Tjark Wendel, Ashcombe Poststelle",
     "Sehr geehrte Frau Voss, zur Forderung APB-0004 übersende ich den heutigen Rücklaufvermerk. Das Schreiben vom 04.09.2026 ging an Rosensteg 8, obwohl im Kundenportal seit dem 19.08.2026 Ahornhof 11 steht. Auf dem Umschlag steht ‚Empfänger unter der angegebenen Anschrift nicht zu ermitteln‘. Der Portalexport zeigt eine Adressänderung, aber keinen Nachweis, wer die Änderung bestätigt hat. In der Mahnstrecke ist das Flag ‚zugestellt‘ durch den nächtlichen Stapel gesetzt worden, weil keine elektronische Fehlermeldung vorlag. Dieses Flag ist kein Postzustellungsnachweis. Ich habe den Brief nicht erneut versandt, weil Ihre Weisung zum Kündigungslauf noch offen ist. Der Ausdruck des Kundenportals ist für die Datenraumabgabe beigefügt; die Unterschrift unter dem alten Formular fehlt in unserer Scanserie. Bitte klären Sie mit dem Erwerber, welche Anschrift für die Überleitungsmitteilung verwendet werden soll.",
     "Ich habe meine Adresse zweimal im Portal geändert und dazu im August telefoniert. Erst meine ehemalige Mitbewohnerin hat mir gestern von einem Bankbrief erzählt. Bitte stellen Sie die Mahnkosten richtig, bevor mir ein anderes Unternehmen schreibt. Meine Mobilnummer hat sich nicht geändert, ich arbeite aber tagsüber am Patienten und kann dann nicht telefonieren."),
    (5, "Ottilie Haberkamm", "Kirschanger 4, 99091 Erfurt", "Rentnerin", "Badumbau", 30000, 0.042, "Titel und Vollstreckung",
     "Abrechnung eines Vollstreckungsauftrags", "Lorenz Tannreuther, Prozessbevollmächtigter der Bank",
     "Sehr geehrte Frau Voss, der vorhandene Vollstreckungsbescheid betrifft nach dem beigefügten Forderungsverzeichnis die im August 2026 bezifferte Teilforderung. Er tituliert nicht automatisch jeden Betrag Ihres Septemberauszugs. Aus dem Auftrag sind am 16.09.2026 180 Euro eingegangen. Meine Kanzlei hat diesen Betrag am 18.09.2026 weitergeleitet; er darf nicht zusätzlich als unmittelbare Kundenzahlung gezählt werden. Das Original des Titels befindet sich noch in meiner Handakte. Die elektronische Kopie ist ohne vollstreckbare Ausfertigung und ohne Zustellungsnachweis in den Datenraum eingestellt worden. Bei einem Forderungsverkauf benötige ich die dokumentierte Rechtsnachfolge und eine Weisung, wer künftig Vollstreckungsaufträge erteilen darf. Die behaupteten weiteren Kosten von 120 Euro stammen aus einer internen Aufwandspauschale der Bank und nicht aus meiner Kostenrechnung. Ich bitte um gesonderte Entscheidung, ob diese Position überhaupt weiterverfolgt werden soll.",
     "Die 180 Euro habe ich an Ihre Kanzlei gezahlt, nicht noch einmal an die Bank. Ich habe nur eine kleine Rente und zahle nach der Absprache im Juli. Es macht mich nervös, dass in Ihrem Brief wieder die gesamte alte Summe steht. Bitte schicken Sie eine Abrechnung, in der ich sehen kann, was mit meinen Zahlungen geschehen ist."),
    (6, "Leander Schlotfeger", "Weidenring 13, 99096 Erfurt", "Koch", "Motorrad und Umzug", 12000, 0.06, "Stundung",
     "Vereinbarung über vorübergehend ausgesetzte Tilgungsleistungen", "Elin Voss, Ashcombe Forderungsservice",
     "Sehr geehrter Herr Schlotfeger, wir bestätigen die am 02.06.2026 besprochene Zahlungserleichterung. Die Tilgungsanteile der Monate Juni, Juli und August 2026 werden bis zum 31.12.2026 gestundet. Die während dieser Monate vertraglich anfallenden Zinsen zahlen Sie jeweils am Monatsende; eine zusätzliche Bearbeitungsgebühr wird nicht erhoben. Ab September 2026 wird wieder die planmäßige Tilgung von 200 Euro zuzüglich der auf den tatsächlichen Kapitalsaldo berechneten Zinsen gezahlt. Über die Verteilung der drei nachzuholenden Tilgungsanteile sprechen wir im November. Eine automatische Fälligstellung des gesamten Darlehens allein wegen der vereinbarten Pause ist nicht vereinbart. Bitte bestätigen Sie die Regelung durch Rücksendung. Unser System kann eine tilgungsfreie Zeit bisher nur als Rückstand anzeigen; dieser technische Status soll im Kundenservice nicht als Mahnfreigabe verwendet werden. Eine Änderung der Sicherheiten haben wir nicht besprochen.",
     "Die Vereinbarung habe ich am 05.06. unterschrieben zurückgegeben. Dass im Portal drei rote Rückstände stehen, verstehe ich nicht. Die Zinsen habe ich bezahlt. Ab September arbeite ich wieder voll. Bitte nennen Sie mir einen Ansprechpartner, der die Vereinbarung kennt, bevor mir der neue Forderungsinhaber eine Kündigung schickt."),
    (7, "Samira Ben Youssef", "Buchenpfad 6, 99099 Erfurt", "Technische Zeichnerin", "Möbel und Waschmaschine", 18000, 0.054, "Widerruf und Formulararchiv",
     "Erklärung zum Widerruf und Bitte um vollständige Vertragsunterlagen", "Samira Ben Youssef",
     "Sehr geehrte Damen und Herren, ich widerrufe meine auf den Abschluss des Darlehens APB-0007 gerichtete Erklärung. In dem mir übersandten PDF fehlt die Seite mit den Angaben zur Widerrufsfrist. Als ich den Kredit abgeschlossen habe, hat die Beraterin mir nur die Seiten 1 bis 3 ausgedruckt. Ich bestreite nicht, die Auszahlung erhalten zu haben; ich bitte um Mitteilung, welche Beträge Sie im Falle einer Rückabwicklung ansetzen. Bitte übersenden Sie mir den vollständigen Datensatz der Dokumentenzustellung, die damalige Formularversion und die Empfangsbestätigung. Der Link in Ihrer Bestätigung führt heute auf ein Formular von 2025, nicht auf meinen Vertrag von 2023. Ich werde die laufenden Beträge zunächst unter Vorbehalt zahlen, um keine weiteren Probleme zu bekommen. Bitte teilen Sie einem Käufer mit, dass diese Frage ungeklärt ist, und stellen Sie meine Nachricht nicht lediglich als allgemeine Beschwerde dar.",
     "Danke für die Antwort. Die angehängte Belehrung von 2025 hilft mir nicht weiter. Ich habe das alte Handy nicht mehr, aber noch die damalige E-Mail. Darin steht nur ‚Ihr Vertrag liegt bereit‘. Können Sie bitte erklären, worauf sich das Häkchen ‚vollständig zugestellt‘ in Ihrem Export stützt? Ich habe keine externe Rechtsvertretung und möchte alles schriftlich."),
    (8, "Gottlieb Birkenseer", "Rosenwinkel 3, 99085 Erfurt", "Elektriker", "Private Wohnungsrenovierung", 27000, 0.048, "Abtretungsklausel",
     "Einwendung gegen einen Wechsel des Forderungsinhabers", "Gottlieb Birkenseer",
     "Sehr geehrte Frau Voss, im September 2023 habe ich mit Herrn Wendel ausdrücklich eine Zusatzvereinbarung unterzeichnet. Darin heißt es, dass eine Abtretung der Ansprüche aus diesem Darlehen an ein nicht zur Bankengruppe gehörendes Unternehmen meiner vorherigen schriftlichen Zustimmung bedarf. Diese Vereinbarung liegt bei meinem Vertrag und ist mit den Unterschriften beider Seiten versehen. Im Datenraum Ihres Verkaufsprojekts steht dagegen offenbar ‚frei übertragbar‘. Ich habe keinem Verkauf zugestimmt. Meine monatlichen Zahlungen laufen weiter; daraus soll keine Zustimmung abgeleitet werden. Bitte übersenden Sie mir den geplanten Namen des Erwerbers und erklären Sie, ob nur die Forderung oder der gesamte Vertrag übertragen werden soll. Ich möchte nicht, dass meine Renovierungsrechnungen an beliebig viele Interessenten verteilt werden. Für eine nachvollziehbare Lösung bin ich erreichbar, aber nicht unter Zeitdruck durch eine pauschale Sammelnachricht.",
     "Den Zusatz habe ich eingescannt. Der Kollege im Service sagte, das sei nur eine unverbindliche Notiz. Auf meiner Seite steht jedoch dieselbe Vertragsnummer und das Datum. Ich habe nichts dagegen, dass Sie mir ein konkretes Angebot machen. Bis dahin zahlen Sie bitte weiterhin auf mein bisheriges Darlehenskonto an und ziehen keine zusätzlichen Gebühren ein."),
    (9, "Nora Chatterjee", "Finkenrain 12, 99097 Erfurt", "Softwareentwicklerin", "Lastenrad und Innenausbau", 24000, 0.06, "Bestrittene Forderung und Aufrechnung",
     "Beanstandung der Doppelabbuchung und Aufrechnungserklärung", "Nora Chatterjee",
     "Sehr geehrte Damen und Herren, im Juni wurde der monatliche Betrag doppelt eingezogen. Sie haben die erste Belastung als Rate und die zweite als freiwillige Sondertilgung verbucht. Ich hatte nur einen Einzug freigegeben und keine Sondertilgung beauftragt. Die zweite Belastung von 456 Euro wurde am 03.07.2026 zurückgegeben; Ihre Liste führt nun eine Rücklastschriftgebühr und einen angeblichen Tilgungsrückstand. Ich rechne vorsorglich mit meinem Rückzahlungsanspruch auf, soweit Sie die doppelte Belastung nicht korrigieren. Den laufenden Kreditvertrag stelle ich im Übrigen nicht infrage. Die Kosten von 25 Euro bestreite ich. Mein Kontoauszug und die Nachricht Ihres Mitarbeiters vom 08.07.2026 liegen Ihnen vor. Bitte geben Sie die Forderung nicht mit dem pauschalen Vermerk ‚unbestritten‘ weiter. Ich erwarte eine Einzelaufstellung, die die Originalbelastung, die doppelte Buchung und die Rückgabe getrennt zeigt.",
     "Sie haben mir nun drei Salden genannt. In der Excel-Liste ist die Rückgabe erfasst, im Septemberbrief fehlt aber die zuvor doppelt erfasste Zahlung. Ich kann ohne einen Buchungsabgleich nicht sagen, welche Zahl stimmen soll. Bitte lassen Sie diese konkrete Frage beantworten und senden Sie nicht wieder die allgemeinen Zahlungsbedingungen."),
    (10, "Hieronymus Fichtelbauer", "Tannensteg 5, 99098 Erfurt", "Rentner", "Dachfenster im selbst genutzten Haus", 36000, 0.042, "Erbfall und Vertretung",
     "Mitteilung des Todesfalls und ungeklärte Erbfolge", "Amelie Fichtelbauer",
     "Sehr geehrte Damen und Herren, mein Vater Hieronymus Fichtelbauer ist am 18.08.2026 verstorben. Ich habe Ihnen eine Kopie der Sterbeurkunde geschickt. Einen Erbschein habe ich nicht. Meine Schwester lebt in Dänemark; wir wissen noch nicht, ob ein Testament existiert. Die Abbuchung im September wurde von der Bank meines Vaters zurückgegeben. Bitte buchen Sie nicht von meinem privaten Konto ab und behandeln Sie meine Nachricht nicht als Schuldübernahme. Ich kümmere mich im Augenblick nur um seine Post. Nach meiner Kenntnis gibt es im Haus eine Mappe mit einer Restschuldversicherung, aber ich habe den Vertrag noch nicht gefunden. Bitte teilen Sie mit, welche Unterlagen Sie tatsächlich benötigen. Eine persönliche Haftungszusage gebe ich nicht. Für weitere Schreiben verwenden Sie zunächst die bisherige Anschrift mit dem Zusatz ‚Nachlass Hieronymus Fichtelbauer‘.",
     "Der Berater hat mich heute als ‚neue Darlehensnehmerin‘ begrüßt. Das stimmt nicht. Ich habe lediglich die Sterbeurkunde übermittelt. Meine Schwester und ich haben die Erbfrage noch nicht geklärt. Wenn Sie die Forderung verkaufen, muss diese Unsicherheit mitgegeben werden. Bitte stellen Sie mir keine Zahlungsvereinbarung mit meinem Namen aus."),
    (11, "Rosalie Ziegenhagen", "Silberweide 2, 99092 Erfurt", "Bibliothekarin", "Solaranlage ohne grundpfandrechtliche Sicherheit", 30000, 0.048, "Sondertilgung und Vorfälligkeitskosten",
     "Ablösungsanfrage und Widerspruch gegen pauschale Kosten", "Rosalie Ziegenhagen",
     "Sehr geehrte Frau Voss, ich möchte den Restkredit im Oktober aus einer Erbschaft ablösen. Bitte erstellen Sie eine Abrechnung zum 20.10.2026. Der mir genannte Betrag von 450 Euro ‚Verkaufskosten‘ ist für mich nicht nachvollziehbar. In meinem Vertrag steht, dass Sondertilgungen ohne gesondertes Entgelt möglich sind. Ich bitte deshalb um Angabe, auf welcher Vereinbarung Ihr zusätzlicher Betrag beruht und wie er berechnet wurde. Bis zu einer geklärten Abrechnung habe ich keine Vollablösung veranlasst. Die Zahlung vom 15.09.2026 in Höhe von 1.000 Euro sollte vollständig auf das Kapital gebucht werden und nicht als Vorauszahlung für Oktober. Bitte bestätigen Sie den neuen Kapitalsaldo. Falls inzwischen ein anderer Gläubiger zuständig wird, benötige ich eindeutige Zahlungsdaten und eine Bestätigung, dass die bisher geleistete Sondertilgung berücksichtigt ist.",
     "Vielen Dank für die neue Aufstellung. Sie enthält weiterhin die 450 Euro, aber keine Berechnung. Im Portal steht meine Sondertilgung noch unter ‚Zahlung ungeklärt‘. Bitte klären Sie zuerst diese beiden Punkte. Ich möchte nicht an einen alten und einen neuen Gläubiger gleichzeitig zahlen."),
    (12, "Cem Krawinkel", "Erlenhain 16, 99086 Erfurt", "Veranstaltungstechniker", "Gebrauchtes Wohnmobil", 42000, 0.054, "Sicherheit und fehlende Urkunde",
     "Unterlagen zur Sicherungsübereignung des Wohnmobils", "Tjark Wendel, Ashcombe Filiale Erfurt",
     "Sehr geehrte Frau Voss, im System ist zur Vertragsnummer APB-0012 eine Sicherungsübereignung des Wohnmobils hinterlegt. Die Zulassungsbescheinigung Teil 2 liegt nach dem alten Tresorverzeichnis in Fach 14. Bei der Inventur vom 28.09.2026 befand sich dort nur eine Kopie; das Original wurde offenbar im Mai zum Halterwechsel herausgegeben. Herr Krawinkel erklärt, das Fahrzeug gehöre inzwischen seiner Ehefrau. Der Vertrag enthält eine Verpflichtung zur Vorlage der Urkunde, aber in der elektronischen Akte fehlt die separat unterschriebene Sicherungsvereinbarung. Wir sollten dem Käufer deshalb nicht allein aufgrund des Systemkennzeichens einen vollständig dokumentierten Sicherheitenbestand bestätigen. Eine Wertermittlung wurde im Juni 2024 vorgenommen; seitdem kennen wir weder Kilometerstand noch Schäden. Bitte entscheiden Sie, ob der Datensatz im Kaufpreisblatt als besichert geführt werden soll, bevor wir die fehlenden Unterlagen erhalten haben.",
     "Das Fahrzeug steht noch bei uns, wird aber von meiner Frau gefahren. Das Originalpapier habe ich nach der Ummeldung im Mai der Filiale gegeben, jedenfalls erinnere ich mich so. Die Kopie schicke ich noch einmal. Ich zahle regelmäßig. Bitte unterstellen Sie mir nicht, ich hätte das Wohnmobil verkauft, nur weil eine Unterlage bei Ihnen nicht auffindbar ist."),
    (13, "Euphemia Dinkelacker", "Haselgasse 8, 99089 Erfurt", "Erzieherin", "Ausbildung der Tochter", 12000, 0.042, "Verbundener Vertrag und Rückabwicklung",
     "Unterrichtsausfall und Einwendungen gegen die weitere Finanzierung", "Euphemia Dinkelacker",
     "Sehr geehrte Damen und Herren, die finanzierte Sprachschule hat seit März keine Leistungen mehr erbracht. Den Kredit habe ich im Schulbüro abgeschlossen; der Berater hatte Ihre Formulare dabei, und der Betrag ging unmittelbar an die Schule. Ich habe die Schule mehrfach zur Wiederaufnahme und anschließend zur Rückzahlung aufgefordert. Die entsprechenden Schreiben finden Sie in meiner Mail vom 17.09.2026. Bitte erläutern Sie, weshalb Sie in Ihrem heutigen Schreiben erklären, die Leistung der Schule gehe die Bank unter keinen Umständen etwas an. Ich habe die Raten bis August zunächst weitergezahlt, weil ich die möglichen Folgen nicht überblicken konnte. Die Septemberzahlung habe ich zurückgestellt. Eine Unterschrift auf Ihrem Entwurf eines Anerkenntnisses werde ich nicht leisten. Bitte prüfen Sie die Vertragsanbahnung und die vorhandenen Vereinbarungen mit der Schule anhand der Unterlagen.",
     "Anbei noch die damalige Begrüßungsnachricht der Schule: ‚Finanzierung über unseren Bankpartner, wir übernehmen die gesamte Abwicklung.‘ Das war für mich ein Paket. Ich habe nie frei über das Darlehensgeld verfügt. Bitte berücksichtigen Sie diese Nachricht auch dann, wenn Ihre Forderung an ein anderes Unternehmen abgegeben wird."),
    (14, "Yara von Lauterbach", "Kornblumenweg 14, 99096 Erfurt", "Grafikdesignerin", "Private Möbelbeschaffung", 18000, 0.06, "Identität und bestrittene Signatur",
     "Bestreiten der elektronischen Vertragserklärung", "Yara von Lauterbach",
     "Sehr geehrte Damen und Herren, ich bestreite, die von Ihnen vorgelegte elektronische Vertragserklärung selbst abgegeben zu haben. Die dort hinterlegte Mobilnummer gehört meinem früheren Partner. Der Betrag ging auf ein Konto, dessen Inhaberin ich nicht bin. Dass meine damalige Anschrift und eine Kopie meines Ausweises verwendet wurden, beweist nicht, dass ich den Vertrag geschlossen habe. Ich habe am 22.09.2026 Anzeige erstattet und Ihnen die Eingangsbestätigung zugesandt. Die früheren Abbuchungen liefen nach meiner Kenntnis von seinem Konto. Bitte bewahren Sie die vollständigen Signatur- und Identifizierungsprotokolle sowie die Daten zur Auszahlungsfreigabe auf. Schicken Sie nicht nur eine Bilddatei mit meinem eingetippten Namen. Ich verlange keine Löschung der Beweise, sondern eine Sperre weiterer automatisierter Forderungsschreiben bis zur Klärung.",
     "Ihr Service antwortet immer mit ‚Sie haben drei Jahre gezahlt‘. Ich habe erläutert, dass die Zahlungen nicht von mir kamen. In Ihrer Exportdatei steht beim Zahlerkonto nur eine interne Nummer. Bitte legen Sie offen, auf welche Belege Sie Ihre Zuordnung stützen. Ich habe keine Ratenvereinbarung bestätigt."),
    (15, "Severin Hopfenmüller", "Kastanienbogen 22, 99091 Erfurt", "Krankenpfleger", "Eigenfinanzierter Umzug", 15000, 0.054, "Dienstleister und Datenzugang",
     "Auskunft über ein unbeabsichtigt offenes Kundenportal", "Severin Hopfenmüller",
     "Sehr geehrte Damen und Herren, über den Link Ihrer Nachricht vom 16.09.2026 konnte ich vorübergehend ein PDF mit dem Namen einer anderen Person sehen. Ich habe die Datei nicht weiterverwendet und den Zugang sofort geschlossen. Bitte teilen Sie mir mit, welche meiner Unterlagen möglicherweise anderen Personen zugänglich waren und wer den Versand veranlasst hat. Ich habe meinen Vertrag mit Ihrer Erfurter Filiale geschlossen. In der technischen Fußzeile der Nachricht steht nun ein Dienstleister in Manchester. Einen Verkauf meines Darlehens kann ich aus den bisherigen Mitteilungen nicht erkennen. Meine Zahlungen laufen pünktlich weiter; ich beanstande den Umgang mit den Daten. Bitte beantworten Sie die konkrete Zugriffsfrage schriftlich und lassen Sie meine Nachricht nicht in der allgemeinen Warteschlange für Mahnungen liegen.",
     "Ich habe inzwischen die Bestätigung, dass das Portal repariert sei. Das beantwortet aber nicht, ob meine Ausweiskopie oder Gehaltsnachweise betroffen waren. Bitte geben Sie dem Erwerber nicht mehr Unterlagen als erforderlich. Mir genügt zuerst eine nachvollziehbare Auskunft und ein Ansprechpartner, der den Vorfall kennt."),
    (16, "Thekla Morgenstern", "Lerchenhain 19, 99085 Erfurt", "Apothekerin", "Private Renovierung", 24000, 0.042, "Gemeinschaftsschuldner und Trennung",
     "Bitte um Entlassung aus dem gemeinsamen Darlehen", "Thekla Morgenstern",
     "Sehr geehrte Frau Voss, mein früherer Ehemann Nils Morgenstern und ich haben das Darlehen gemeinsam aufgenommen. Seit der Trennung im April zahlt er nach unserer privaten Vereinbarung die Raten. Ich bitte um schriftliche Prüfung, ob Sie mich aus dem Vertrag entlassen. Eine Einigung zwischen ihm und mir ist keine Bestätigung der Bank, das ist uns bewusst. Ihr Kundenservice hat mich dennoch aus der Verteilerliste gestrichen und in der Übersicht nur noch Nils aufgeführt. Bitte berichtigen Sie diese Darstellung, solange Sie keine Entlassung erklärt haben. Ich möchte Kopien aller Änderungen erhalten. Die neue Anschrift meines früheren Ehemanns kenne ich nicht sicher; die letzte Zahlung kam von demselben Konto wie zuvor. Sollte das Darlehen verkauft werden, muss der Käufer den gemeinsamen Vertrag und nicht nur einen verkürzten Personendatensatz erhalten.",
     "Es gibt keine von mir unterschriebene Vertragsübernahme. Ich habe nur gefragt, ob das möglich ist. Bitte schicken Sie ein tatsächliches Angebot und nennen Sie die Voraussetzungen. Mein ehemaliger Mann soll nicht erneut allein für uns beide unterschreiben. Den Kaufpreis Ihres Portfolios betrifft unsere private Vereinbarung nicht."),
    (17, "Abdul Rahman Eberlein", "Mühlenbogen 10, 99094 Erfurt", "Straßenbahnfahrer", "Gebrauchtwagen und Haushaltsgeräte", 21000, 0.048, "Zahlungszuordnung und Stichtag",
     "Ungeklärter Zahlungseingang im Sammelkonto", "Alma Wurzel, Ashcombe Buchhaltung",
     "Sehr geehrte Frau Voss, der Eingang von 720 Euro am 30.09.2026 um 16.42 Uhr steht im Hauptbuch auf dem Sammelkonto 1690. Als Verwendungszweck ist ‚Eberlein Kredit September Oktober‘ angegeben. Die operative Darlehensliste wurde um 12.00 Uhr geschlossen und enthält den Betrag noch nicht. Ich habe den Eingang nach Abgleich der Zahlerdaten APB-0017 zugeordnet; die technische Buchung erfolgte am 01.10.2026 mit Wertstellung 30.09.2026. Bitte führen Sie eine Stichtagsbrücke und verändern Sie den unveränderten Mittagsauszug nicht rückwirkend. Der Verkauf sollte eindeutig regeln, wem dieser Betrag zusteht und wer eine Doppelabbuchung verhindert. Die Bankabstimmung stimmt nur dann, wenn das Sammelkonto und das Darlehensnebenbuch gemeinsam betrachtet werden. Im Septemberabschluss ist die Position noch als unzugeordnete Zahlung ausgewiesen.",
     "Ich habe die 720 Euro am Monatsletzten überwiesen und den Kontoauszug angehängt. Das sollte den Rückstand aus September ausgleichen und den Rest auf das Kapital bringen. Bitte rechnen Sie nicht so, als sei die Zahlung erst im Oktober bei Ihnen eingegangen. Eine weitere Abbuchung über den gesamten alten Betrag wäre für mich ein echtes Problem."),
    (18, "Wendelin Nussbaumer", "Apfelgarten 23, 99099 Erfurt", "Feinmechaniker", "Haushaltsgeräte und Umzug", 18000, 0.054, "Ordnungsgemäßer Verlauf",
     "Saldenbestätigung auf Kundenanfrage", "Elin Voss, Ashcombe Forderungsservice",
     "Sehr geehrter Herr Nussbaumer, wir bestätigen Ihnen den Darlehenskontostand gemäß der beigefügten Aufstellung zum 30.09.2026. Alle bis dahin vereinbarten Monatszahlungen sind eingegangen. Aus der Saldenbestätigung ergibt sich weder eine vorzeitige Fälligstellung noch eine Änderung Ihrer Zahlungsbedingungen. Sie können den beiliegenden Buchungsauszug für Ihre private Übersicht verwenden. Für einen möglichen Wechsel des Forderungsinhabers haben wir noch keine Zahlungsumstellung veranlasst. Bis zu einer gesonderten nachvollziehbaren Mitteilung bleiben die Ihnen bekannten Zahlungswege bestehen. Bitte prüfen Sie neue Kontodaten bei Zweifeln über die bisher bekannte Kontaktadresse. Eine Anfrage nach Ihrer Portal-PIN wird unsere Filiale im Zuge des Verkaufs nicht stellen. Über eine von Ihnen gewünschte Sondertilgung können wir unabhängig vom Verkaufsprojekt sprechen.",
     "Vielen Dank, die Übersicht passt zu meinen Unterlagen. Ich brauche eigentlich nur eine Bestätigung, dass sich die vereinbarte monatliche Tilgung nicht durch den Verkauf ändert. Bitte lassen Sie mir rechtzeitig eine Nachricht zukommen, falls ich den Dauerauftrag anpassen muss. Ich möchte nicht auf eine gefälschte Zahlungsaufforderung hereinfallen."),
    (19, "Lotte Izumi Greifenstein", "Schlehenpfad 17, 99097 Erfurt", "Logopädin", "Private Einrichtung eines Arbeitszimmers", 30000, 0.06, "Vergleich und bedingter Erlass",
     "Bestätigung einer Vergleichsvereinbarung", "Elin Voss, Ashcombe Forderungsservice",
     "Sehr geehrte Frau Greifenstein, wir bestätigen den am 10.09.2026 schriftlich vereinbarten Vergleich. Sie zahlen bis zum 15.10.2026 einen Betrag von 8.000 Euro. Mit vollständigem fristgerechtem Eingang erlassen wir die danach verbleibende Forderung aus APB-0019 einschließlich der bis dahin von uns geführten Nebenforderungen. Bis zur Entscheidung über die rechtzeitige Erfüllung setzen wir weitere Einziehungsmaßnahmen aus. Die Vereinbarung enthält keine automatische Wiederauflebensklausel für bereits endgültig erlassene Beträge. Bei nicht fristgerechtem Eingang prüfen wir das weitere Vorgehen auf Grundlage des vollständigen Textes und Ihrer Umstände. Der Verkauf unserer Forderung darf die zugesagte Vergleichswirkung nicht beseitigen. Bitte verwenden Sie im Verwendungszweck die Vertragsnummer und das Wort Vergleich. Eine Zahlung von 2.000 Euro am 25.09.2026 haben wir als erste Teilzahlung erhalten; die weitere Zahlung steht noch aus.",
     "Die ersten 2.000 Euro sind überwiesen, die übrigen 6.000 Euro kommen nach dem Verkauf meines Autos bis zum 15.10. Ich habe dem Vergleich zugestimmt, weil dann endlich alles abgeschlossen sein soll. Bitte bestätigen Sie, dass ein Käufer den Erlass nicht später zurücknehmen kann. Ich kann nicht mehr bezahlen als vereinbart."),
    (20, "Milo Wackernagel", "Eschenhof 15, 99086 Erfurt", "Fahrradmechaniker", "Private Weiterbildung", 12000, 0.048, "Altunterlagen und Zahlungsnachweis",
     "Nachreichung der unterschriebenen Vertragsfassung", "Tjark Wendel, Ashcombe Filiale Erfurt",
     "Sehr geehrte Frau Voss, zur Nummer APB-0020 habe ich das Papierarchiv nochmals durchsucht. Der im Datenraum gespeicherte Vertrag ist die vor der Unterschrift versandte Fassung; die letzte Seite zeigt noch keine Unterschrift. Im Archiv findet sich dagegen eine beidseitig unterschriebene Fassung mit demselben Nettobetrag und derselben Laufzeit, aber einer anderen E-Mail-Adresse. Ich habe die vollständige Fassung heute bereitgestellt. Für das Datum der Aushändigung haben wir keinen gesonderten Empfangsbeleg; im Kundenservice steht nur ‚Mappe mitgegeben‘. Die Auszahlung an Herrn Wackernagel ist über die Buchungsreferenz AZ-APB-0020 nachvollziehbar. Er hat die Auszahlung nie bestritten und zahlt regelmäßig. Bitte führen Sie beide Fassungen mit eindeutiger Versionsbezeichnung, damit der Erwerber nicht versehentlich einen unsignierten Entwurf als Originalbestand behandelt.",
     "Ich habe den unterschriebenen Vertrag noch zu Hause. Auf meiner Kopie steht eine handschriftliche Telefonnummer des damaligen Beraters. Die E-Mail-Adresse wurde später geändert, weil ich keinen Zugang mehr zum alten Postfach hatte. Wenn die neue Bank mir schreibt, soll sie bitte die aktuelle Adresse verwenden. Ich schicke Ihnen die Titelseite und die letzte Seite noch einmal."),
]


def build_model():
    cases = []
    for seed in SEEDS:
        n, name, address, occupation, purpose, amount, rate, topic, heading, sender, letter, reply = seed
        initial = amount * 100
        amort = initial // 60
        # Die historische Saldenübernahme ist eine Verkäuferangabe, keine Aktenprüfung.
        paid_before = 4 if n == 2 else 24
        opening_p = initial - amort * paid_before
        opening_i = 7200 if n == 2 else 0
        p, interest, costs = opening_p, opening_i, 0
        rows = []
        total_principal_paid = total_interest_paid = total_costs_paid = 0
        for m in range(12):
            year, month = (2025, m + 10) if m < 3 else (2026, m - 2)
            date = f"{year}-{month:02d}-{calendar.monthrange(year, month)[1]}"
            charged = int((Decimal(p) * Decimal(str(rate)) / 12).quantize(Decimal(1), rounding=ROUND_HALF_UP))
            pay_p, pay_i, add_cost = amort, charged, 0
            note = "Monatszahlung gemäß Zahlungsauftrag"
            if n == 1 and m >= 8: pay_p = pay_i = 0; note = "Kein Eingang seit Verfahrenseröffnung"
            if n == 2: pay_p = pay_i = 0; note = "Altbestandsbuchung; keine Kundenzahlung"
            if n in (3,4) and m >= 9: pay_p = pay_i = 0; note = "Kein zugeordneter Monatseingang"
            if n == 5 and m >= 7: pay_p, pay_i = (18000,0) if m == 11 else (0,0); note = "Weiterleitung Kanzlei" if m == 11 else "Vollstreckungsakte; kein Eingang"
            if n == 6 and m in (8,9,10): pay_p = 0; note = "Tilgung nach Vereinbarung ausgesetzt; Zins bezahlt"
            if n == 10 and m >= 10: pay_p = pay_i = 0; note = "Todesfall; kein Eingang"
            if n in (13,14,17) and m == 11: pay_p = pay_i = 0; note = "Kein Eingang im Mittagsauszug"
            if n == 19 and m >= 7: pay_p = pay_i = 0; note = "Vergleichsverhandlung"
            if n == 19 and m == 11: pay_p = 200000; pay_i = 0; note = "Vergleichsteilzahlung ausdrücklich auf Kapital"
            if n == 11 and m == 11: pay_p += 100000; note = "Monatszahlung plus 1000 EUR Sondertilgung"
            if n == 1 and m == 7: add_cost = 3600
            if n == 5 and m == 10: add_cost = 12000
            if n == 9 and m == 9: add_cost = 2500
            if n == 11 and m == 11: add_cost = 45000
            if n in (3,4) and m in (10,11): add_cost = 600
            interest += charged - pay_i
            costs += add_cost
            p -= pay_p
            total_principal_paid += pay_p; total_interest_paid += pay_i
            rows.append(dict(date=date, reference=f"J-{n:04d}-{m+1:02d}", principal_start=p+pay_p,
                             interest_charge=charged, costs_charge=add_cost, paid_principal=pay_p,
                             paid_interest=pay_i, paid_costs=0, cash=pay_p+pay_i, principal_end=p,
                             interest_end=interest, costs_end=costs, total_end=p+interest+costs,note=note))
        # Vertraglicher Monatszins: jeden Betrag auf Cent runden, erst dann summieren.
        contract_cash = [amort + int((Decimal(initial-amort*k)*Decimal(str(rate))/12).quantize(Decimal(1),rounding=ROUND_HALF_UP)) for k in range(60)]
        total_contract = sum(contract_cash)
        lo, hi = 0.0, 0.05
        for _ in range(100):
            mid=(lo+hi)/2
            present=sum(value/(1+mid)**(k+1) for k,value in enumerate(contract_cash))
            if present > initial:lo=mid
            else:hi=mid
        effective=(1+(lo+hi)/2)**12-1
        cases.append(dict(id=f"APB-{n:04d}", number=n, name=name, address=address,
            occupation=occupation, purpose=purpose, initial=initial, annual_rate=rate,
            effective_rate=effective, monthly_principal=amort, contract_date="2023-09-25",
            disbursement_date="2023-09-30", first_payment="2023-10-31", last_payment="2028-09-30",
            scheduled_total=total_contract, scheduled_payments=contract_cash, topic=topic, heading=heading, sender=sender,
            letter=letter, reply=reply, opening_principal=opening_p, opening_interest=opening_i,
            principal=p, interest=interest, costs=costs, balance=p+interest+costs,
            payments_principal=total_principal_paid,payments_interest=total_interest_paid,
            payments_costs=total_costs_paid, journal=rows))
    portfolio = []
    for n in range(1,1001):
        c = cases[n-1] if n <= 20 else None
        principal = c['principal'] if c else (6000 + (n*137)%24000) * 100
        interest = c['interest'] if c else (n%9)*650
        costs = c['costs'] if c else (1500 if n%29 == 0 else 0)
        # Verkäuferseitiger Preisparameter, nicht Risiko-Rating des DD-Teams.
        price_ratio = 0.92 if n > 20 else {1:.24,2:.08,3:.65,5:.42,14:.10,19:.55}.get(n,.90)
        portfolio.append(dict(id=f"APB-{n:04d}",name=c['name'] if c else f"Bestandskunde {n:04d}",
            scope="Detailakte" if c else "Nur Verkäuferexport",principal=principal,
            interest=interest,costs=costs,balance=principal+interest+costs,
            seller_price_ratio=price_ratio,price=round(principal*price_ratio),
            status=c['topic'] if c else ("Zahlungslücke im Export" if n%29==0 else "Regelbestand laut Verkäufer"),
            source=f"{c['id']}_02_Kontojournal.pdf" if c else "Export KREDIT-ERF-20260930, Einzelbelege nicht im Datenraum"))
    totals = {k:sum(r[k] for r in portfolio) for k in ('principal','interest','costs','balance','price')}
    return dict(schema_version=1,case_slug=SLUG,as_of="2026-09-30",working_date="2026-10-10",
        seller=BANK,buyer=BUYER,cases=cases,portfolio=portfolio,totals=totals,
        adjustments=[dict(id="BR-01",loan="APB-0017",amount=-72000,kind="Sammelkonto 1690",value_date="2026-09-30",booking_date="2026-10-01",source="APB-0017_03_Schreiben.pdf")],
        general_ledger_net=totals['balance']-72000,
        branch_accounts=dict(impairment_individual=38500000,impairment_portfolio=30000000,
            book_value=totals['balance']-72000-68500000,
            income_2025=[124000000,-42000000,-9600000,-14500000,-8500000,-16500000,-8200000],
            income_2026_ytd=[94500000,-32700000,-7200000,-13200000,-10100000,-23100000,-6900000],
            categories=['Zinserträge','Personal','Miete und Betrieb','IT und Systeme','Externer Service','Risikovorsorge','Zentrale Umlage']),
        source_note="Verkäuferexport und Buchungsannahmen der fiktiven Akte; keine DD-Feststellung. 20 Fälle bewusst ausgewählt, keine Zufallsstichprobe.",
        legal_sources_checked=[dict(date="2026-10-10",url=url) for url in [
            "https://www.gesetze-im-internet.de/bgb/__355.html",
            "https://www.gesetze-im-internet.de/bgb/__356b.html",
            "https://www.gesetze-im-internet.de/bgb/__357b.html",
            "https://www.gesetze-im-internet.de/bgbeg/art_247__6.html"]])


def stable_docx(path):
    with zipfile.ZipFile(path) as z:
        data={n:z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
        for name in sorted(data):
            item=zipfile.ZipInfo(name,(2026,10,10,12,0,0)); item.compress_type=zipfile.ZIP_DEFLATED
            z.writestr(item,data[name])


def docx(path,title,intro,sections,meta=""):
    d=Document(); sec=d.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.1); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(2.4);sec.right_margin=Cm(2.2)
    for sn in ('Normal','Title','Heading 1','Heading 2'):
        st=d.styles[sn];st.font.name='Times New Roman';st.font.size=Pt(11 if sn=='Normal' else 15 if sn=='Title' else 12)
        st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.space_after=Pt(7)
        fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+key),None)
        for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Times New Roman')
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.styles['Normal'].paragraph_format.line_spacing=1.08
    d.core_properties.created=datetime(2026,10,10,12,0,0,tzinfo=timezone.utc)
    d.core_properties.modified=d.core_properties.created;d.core_properties.author="Ashcombe Projektteam Erfurt"
    if meta: d.add_paragraph(meta)
    d.add_paragraph(title,'Title');d.add_paragraph(intro)
    for heading,paras in sections:
        d.add_heading(heading,1)
        for para in paras if isinstance(paras,list) else [paras]:d.add_paragraph(para)
    foot=sec.footer.paragraphs[0];foot.alignment=2
    foot.add_run("Projekt Lindenbogen | Seite ")
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
    for element in list(d.element.iter(qn('w:pBdr'))):element.getparent().remove(element)
    for style in d.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.save(path);stable_docx(path)


STYLES=getSampleStyleSheet()
STYLES.add(ParagraphStyle(name='BankBody',fontName='Times-Roman',fontSize=11,leading=14,spaceAfter=8))
STYLES.add(ParagraphStyle(name='BankTitle',fontName='Times-Bold',fontSize=16,leading=20,spaceAfter=14))
STYLES.add(ParagraphStyle(name='BankHead',fontName='Times-Bold',fontSize=12,leading=15,spaceBefore=10,spaceAfter=8))
STYLES.add(ParagraphStyle(name='BankCell',fontName='Times-Roman',fontSize=9,leading=11,spaceAfter=0))


def par(text,style='BankBody'):
    return Paragraph(escape(str(text)).replace('\n','<br/>'),STYLES[style])


def pdf(path,title,blocks,meta='',table=None):
    story=[]
    if meta:story.append(par(meta))
    story.append(par(title,'BankTitle'))
    for b in blocks:
        if isinstance(b,tuple):story.append(par(b[0],'BankHead'));story.extend(par(p) for p in b[1])
        else:story.append(par(b))
    if table:
        heads,rows,widths=table
        cells=[[par(x,'BankCell') for x in heads]]+[[par(x,'BankCell') for x in row] for row in rows]
        t=Table(cells,colWidths=[w*mm for w in widths],repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DDE5EA')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#CDD3D7')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        story.append(t)
    def page(canvas,doc):
        canvas.setFont('Times-Roman',9);canvas.drawString(24*mm,14*mm,'Projekt Lindenbogen | interne Unterlage');canvas.drawRightString(187*mm,14*mm,f'Seite {doc.page}')
    SimpleDocTemplate(str(path),pagesize=(210*mm,297*mm),leftMargin=24*mm,rightMargin=23*mm,topMargin=20*mm,bottomMargin=23*mm,title=title,author='Projektteam Lindenbogen',invariant=1).build(story,onFirstPage=page,onLaterPages=page)


def eml(path,subject,sender,to,body,day=8,attachment=None):
    e=EmailMessage(policy=SMTP);e['From']=sender;e['To']=to;e['Subject']=subject
    e['Date']=f"{['Mon','Tue','Wed','Thu','Fri','Sat','Sun'][datetime(2026,10,day).weekday()]}, {day:02d} Oct 2026 09:24:00 +0200"
    e['Message-ID']=f"<{path.stem.lower()}@ashcombe.example.invalid>"
    e.set_content(body)
    if attachment:
        e.add_attachment(attachment.read_bytes(),maintype='application',subtype='pdf',filename=attachment.name)
        e.set_boundary('=_lindenbogen_'+path.stem)
    path.write_bytes(e.as_bytes())


def contract(c):
    extra="Es werden keine Sicherheiten bestellt. Ein Zugriff auf Arbeitseinkommen ist nicht vereinbart."
    if c['number']==12: extra="Der Darlehensnehmer verpflichtet sich, das finanzierte Wohnmobil nach Maßgabe einer gesondert zu unterzeichnenden Sicherungsvereinbarung sicherungsweise zu übereignen. Diese Verpflichtung ersetzt die gesonderte dingliche Einigung nicht. Die Originalurkunde der Zulassungsbescheinigung Teil 2 soll bei der Filiale verwahrt werden."
    transfer="Die Bank darf Forderungen aus diesem Vertrag im Rahmen der gesetzlichen und vertraglichen Voraussetzungen übertragen. Die bestehenden Einwendungen und vereinbarten Konditionen werden durch einen Gläubigerwechsel nicht einseitig geändert. Eine Übertragung sämtlicher Vertragspflichten ist mit dieser Erklärung nicht vereinbart."
    if c['number']==8:transfer="Abweichend vom Formular bedarf die Abtretung der Ansprüche aus diesem Darlehen an ein nicht zur Bankengruppe gehörendes Unternehmen der vorherigen schriftlichen Zustimmung des Darlehensnehmers. Eine bloße Ankündigung ersetzt die Zustimmung nicht. Diese individuell besprochene Zusatzabrede geht entgegenstehenden Formularbestimmungen vor."
    borrower=c['name']+(", gemeinsam mit Nils Morgenstern, gleiche damalige Anschrift" if c['number']==16 else "")
    total=c['scheduled_total']
    sections=[
      ('1 Vertragsparteien und Darlehenszweck',[f"Darlehensgeberin ist {BANK}, {ADDR}. Die Erfurter Niederlassung handelt für die englische Gesellschaft und ist keine eigenständige juristische Person. Ansprechpartner für diesen Vertrag ist Tjark Wendel, service@ashcombe.example.invalid. Darlehensnehmer ist {borrower}, {c['address']}. Der angegebene Beruf lautet {c['occupation']}.",f"Das Darlehen dient folgendem privaten Zweck: {c['purpose']}. Eine grundpfandrechtliche Besicherung wird nicht vereinbart. Die Nutzung als laufende Geschäftskreditlinie ist ausgeschlossen. Der Darlehensnehmer erklärt seine Angaben zu Einkommen und bestehenden Zahlungspflichten nach bestem Wissen; die Bank hat die eingereichten Unterlagen vor ihrer Kreditentscheidung gesondert zu prüfen."]),
      ('2 Auszahlung und Nettodarlehensbetrag',[f"Der Nettodarlehensbetrag beträgt {money(c['initial'])}. Die Auszahlung erfolgt am 30.09.2023 in einer Summe nach beiderseitiger Unterzeichnung und Freigabe der angegebenen Empfangsverbindung. Die Auszahlungsreferenz lautet AZ-{c['id']}. Es werden keine Vermittlungs- oder Abschlusskosten einbehalten. Eine Kreditrestschuldversicherung ist weder Voraussetzung noch Bestandteil dieses Darlehens.","Der Darlehensnehmer kann ausbleibende oder abweichende Auszahlung schriftlich beanstanden. Eine Auszahlung an Dritte bedarf eines gesonderten dokumentierten Zahlungsauftrags. Bei APB-0013 ist als Empfänger die Sprachschule benannt; der gesonderte Auftrag ist zusammen mit den Schulunterlagen aufzubewahren." if c['number']==13 else "Der Darlehensnehmer kann ausbleibende oder abweichende Auszahlung schriftlich beanstanden. Eine Auszahlung an Dritte bedarf eines gesonderten dokumentierten Zahlungsauftrags. Eine interne Buchung ersetzt den tatsächlichen Auszahlungsnachweis nicht."]),
      ('3 Laufzeit und Tilgung',[f"Die Laufzeit beginnt mit der Auszahlung und endet bei vertragsgemäßer Durchführung am 30.09.2028. Der Darlehensnehmer leistet 60 monatliche Tilgungsanteile von jeweils {money(c['monthly_principal'])}; der erste ist am 31.10.2023, der letzte am 30.09.2028 fällig. Hinzu kommen die für den jeweiligen Monat nach Nummer 4 berechneten Zinsen. Die monatliche Gesamtzahlung sinkt deshalb bei planmäßigem Verlauf. Es handelt sich nicht um eine gleichbleibende Annuität.","Die Zahlungen werden jeweils am letzten Kalendertag des Monats fällig. Zahlungsweg und Einzugsermächtigung werden gesondert vereinbart. Der Darlehensnehmer kann jederzeit einen Tilgungsplan mit den künftigen Fälligkeiten und der Aufteilung in Kapital und Zinsen verlangen. Abweichende Stundungen oder Zahlpausen werden mit Bezug auf die Vertragsnummer dokumentiert; ein Systemkennzeichen allein verändert den Vertrag nicht."]),
      ('4 Zinsen und Gesamtbetrag',[f"Der gebundene Sollzinssatz beträgt für die gesamte vereinbarte Laufzeit {c['annual_rate']*100:.2f} Prozent jährlich. Bei monatlicher Zahlung beträgt der Zins für einen vollen Vertragsmonat ein Zwölftel dieses Satzes, angewandt auf den Kapitalsaldo zu Monatsbeginn. Einzelne Zinsbeträge werden kaufmännisch auf Cent gerundet. Der effektive Jahreszins beträgt bei planmäßiger Auszahlung und Rückzahlung {c['effective_rate']*100:.2f} Prozent. Berechnungsannahmen sind vollständige Auszahlung am 30.09.2023, 60 Monatszahlungen und das Fehlen zusätzlicher Pflichtkosten.",f"Bei unverändert planmäßiger Durchführung beträgt der Gesamtbetrag aller Zahlungen {money(total)}. Davon entfallen {money(total-c['initial'])} auf Zinsen. Die erste Zahlung beträgt {money(c['monthly_principal']+cent(c['initial']/100*c['annual_rate']/12))}. Dieser Gesamtbetrag ist kein jederzeit fälliger Ablösebetrag. Sonderzahlungen, Stundungen und tatsächliche Abweichungen vom Zahlungsplan können den späteren Zahlungsstrom verändern."]),
      ('5 Zahlungszuordnung und Kontoauszüge',["Die Bank weist in Kontoauszügen Kapital, Zinsen, Kosten und eingegangene Zahlungen getrennt aus. Beanstandete Kosten werden als eigene Position kenntlich gemacht. Eine Kundenzahlung wird nicht allein wegen eines Verkaufs oder eines Servicerwechsels nochmals verlangt. Für Zahlungen mit ausdrücklicher Tilgungsbestimmung ist deren Wirksamkeit zu prüfen; gesetzliche Zuordnungsregeln bleiben unberührt.","Der Darlehensnehmer kann Buchungsfehler unter Angabe des Zahlungsdatums und eines Belegs melden. Die Bank prüft auch Zahlungseingänge auf Sammelkonten. Eine Saldenmitteilung enthält kein Anerkenntnis des Kunden und ersetzt weder eine Fälligkeitsprüfung noch einen Vollstreckungstitel."]),
      ('6 Sondertilgung und Sicherheiten',["Freiwillige Sondertilgungen sind jederzeit ohne gesondertes Entgelt möglich. Sie vermindern nach Eingang den Kapitalsaldo und damit die künftig hierauf anfallenden Zinsen. Die Bank erstellt auf Anfrage eine Abrechnung zu einem benannten Ablösedatum. Unbekannte spätere Zahlungen oder Kosten werden nicht als bereits feststehende Position in eine Abrechnung aufgenommen.",extra]),
      ('7 Zahlungsstörung und Beendigung',["Bei Zahlungsschwierigkeiten soll der Darlehensnehmer frühzeitig mit der Filiale Kontakt aufnehmen. Eine Nachfristsetzung oder Kündigung muss die einschlägigen gesetzlichen Voraussetzungen erfüllen. Weder das bloße Vorliegen zweier unbezahlter Buchungspositionen noch ein Verkauf der Forderung genügt für sich genommen als Kündigungsgrund. Die Bank dokumentiert Zugang und Inhalt ihrer Erklärungen sowie angebotene Lösungsgespräche.","Weitere Zinsen und Kosten bei Verzug werden nur nach den tatsächlich einschlägigen gesetzlichen und vertraglichen Voraussetzungen verlangt und nachvollziehbar berechnet. Eine interne Aufwandspauschale wird mit dieser Klausel nicht vereinbart. Eine titulierte Teilforderung wird nicht ohne gesonderte Grundlage auf andere Forderungsbestandteile ausgeweitet."]),
      ('8 Informationen zum Widerruf',["Der Darlehensnehmer kann seine Vertragserklärung innerhalb von 14 Tagen ohne Angabe von Gründen widerrufen. Der Lauf der Widerrufsfrist setzt den Vertragsschluss und den Erhalt der gesetzlich erforderlichen Vertragsunterlagen und Pflichtangaben voraus. Die vollständige, dem Vertrag zugeordnete Widerrufsinformation wird auf einem dauerhaften Datenträger zusammen mit der Vertragsabschrift übergeben. Zur Wahrung der Frist genügt die rechtzeitige Absendung des Widerrufs, wenn die Erklärung auf einem dauerhaften Datenträger erfolgt.",f"Der Widerruf ist an {BANK}, {ADDR}, widerruf@ashcombe.example.invalid, zu richten. Nach einem wirksamen Widerruf sind empfangene Leistungen nach Maßgabe der anwendbaren Vorschriften zurückzugewähren. Für den Zeitraum zwischen Auszahlung und Rückzahlung wird bei vollständiger Auszahlung ein täglicher Zinsbetrag von {money(round(c['initial']*c['annual_rate']/365))} ausgewiesen. Die Rückzahlung hat spätestens innerhalb von 30 Tagen nach Absendung der Widerrufserklärung zu erfolgen. Gesetzliche Rechte aus einem verbundenen Vertrag werden durch diese Darstellung nicht ausgeschlossen."]),
      ('9 Übertragung und Kommunikation',[transfer,"Die Bank informiert über einen Wechsel des Forderungsinhabers mit eindeutiger Bezeichnung des neuen Gläubigers und einem überprüfbaren Zahlungsweg. Der Darlehensnehmer darf bei widersprüchlichen Mitteilungen einen Nachweis verlangen. Zugangsdaten oder eine Portal-PIN werden für die Überleitungsmitteilung nicht angefordert. Gesetzliche Informations- und Datenschutzpflichten werden gesondert erfüllt."]),
      ('10 Beschwerden und Vertragsunterlagen',["Beschwerden können an beschwerde@ashcombe.example.invalid oder schriftlich an die Erfurter Filiale gerichtet werden. Die Bank antwortet unter Bezug auf die konkrete Vertragsnummer. Eine Beschwerde ersetzt keine fristgebundene Erklärung. Zuständige Aufsichts- und außergerichtliche Streitbeilegungsstellen sowie die Zugangsvoraussetzungen ergeben sich aus dem gesondert übergebenen Informationsblatt in der damals geltenden Fassung; dieses Blatt ist Bestandteil der ausgegebenen Vertragsmappe.","Änderungen werden so dokumentiert, dass ihr Inhalt, ihre Beteiligten und ihr Zeitpunkt nachvollziehbar bleiben. Der Kunde erhält eine Vertragsabschrift. Die Parteien erklären mit ihrer Unterzeichnung nicht den Zugang von Unterlagen, die ihnen tatsächlich nicht ausgehändigt oder übermittelt wurden."]),
      ('11 Unterzeichnung und Ausgabe',[f"Erfurt, 25.09.2023. Für die Bank: Tjark Wendel. Darlehensnehmer: {borrower}. Die archivierte Abschrift gibt die Namen der Unterzeichnenden als Text wieder. Die Originalsignaturen und gegebenenfalls elektronischen Signaturprotokolle sind gesondert aufzubewahren.","Archivvermerk: Im elektronischen Versandpaket ist die Widerrufsseite nicht enthalten; die Filiale behauptet eine Ausgabe im Papierpaket. Ein separater Empfangsbeleg fehlt." if c['number']==7 else "Archivvermerk: Das für APB-0014 gespeicherte elektronische Signaturprotokoll weist eine vom Kunden bestrittene Mobilnummer aus; die bloße Namenswiedergabe in dieser Abschrift klärt die Signaturfrage nicht." if c['number']==14 else "Archivvermerk: Diese Fassung ersetzt den zuvor in den Datenraum eingestellten unsignierten Entwurf. Ein gesonderter Aushändigungsbeleg liegt nicht bei." if c['number']==20 else "Archivvermerk: Die Vertragsabschrift wird mit der Vertragsnummer geführt; Hinweise auf gesonderte Anlagen belegen deren tatsächlichen Zugang nicht."])
    ]
    return sections


def write_cases(model):
    for c in model['cases']:
        stem=c['id']; title=f"Verbraucherdarlehensvertrag {stem}"
        docx(OUT/f'{stem}_01_Darlehensvertrag.docx',title,
            f"Vertrag zwischen {BANK} und {c['name']} über {money(c['initial'])}. Vertragsabschrift aus dem Erfurter Bestandsarchiv; vereinbarter Zahlungsbeginn 31.10.2023.",contract(c),BANK+'\n'+ADDR)
        rows=[]
        for j in c['journal']:
            rows.append([j['date'],money(j['interest_charge']),money(j['cash']),money(j['principal_end']),money(j['interest_end']+j['costs_end'])])
        pdf(OUT/f'{stem}_02_Kontojournal.pdf',f'Darlehenskonto {stem}',[
            f"Kontoinhaber: {c['name']}, {c['address']}. Auszug der Buchungsdaten vom 01.10.2025 bis 30.09.2026, Stand des Exports 30.09.2026, 12.00 Uhr. Alle Beträge in EUR.",
            f"Übernommener Saldo am 01.10.2025: Kapital {money(c['opening_principal'])}; Zinsen {money(c['opening_interest'])}; Kosten 0,00 EUR. Diese Salden stammen aus dem Vorjahresexport. Einzelbuchungen vor Oktober 2025 sind nicht Bestandteil dieses Auszugs.",
            f"Zum Stichtag weist das Nebenbuch Kapital von {money(c['principal'])}, gebuchte unbezahlte Zinsen von {money(c['interest'])} und Kosten von {money(c['costs'])} aus. Buchungssaldo: {money(c['balance'])}. Im Auszugszeitraum wurden {money(c['payments_principal'])} auf Kapital und {money(c['payments_interest'])} auf Zinsen gebucht.",
            "Die Buchungsdarstellung enthält keine Aussage über Fälligkeit, Einwendungen, Durchsetzbarkeit oder einen zulässigen Einziehungsbetrag. Zinsen werden im Altsystem auch nach Statusänderungen weitergebucht. Zahlungen nach dem Exportzeitpunkt stehen gegebenenfalls auf dem Sammelkonto. Die vollständigen Zuordnungsschlüssel enthält die Arbeitsmappe.",
        ],BANK,(['Datum','Zinsbuchung','Zahlung','Kapitalende','Zinsen und Kosten'],rows,[26,29,32,36,40]))
        pdf(OUT/f'{stem}_03_Schreiben.pdf',c['heading'],[
            f"Vorgang {stem} | {c['name']} | 30.09.2026",c['letter'],
            "Bitte führen Sie den gesamten Schriftwechsel unter der genannten Vertragsnummer. Eine Weitergabe an das Projektteam soll die Einwendungen und die offenen Nachweise vollständig enthalten.",
            'Mit freundlichen Grüßen\n'+c['sender']
        ],c['sender']+'\n'+(c['address'] if c['sender']==c['name'] else ADDR))
        correspondent='Amelie Fichtelbauer' if c['number']==10 else c['name']
        contract_wording='eine Vertragsänderung' if c['number']==10 else 'eine Änderung meines Vertrags'
        email= re.sub(r'[^a-z0-9]','',correspondent.lower().replace('ü','ue').replace('ö','oe').replace('ä','ae'))+'@kunden.example.invalid'
        body=(f"Sehr geehrte Frau Voss,\n\nzu {stem} antworte ich auf Ihr Schreiben und die Datenanfrage. {c['reply']}\n\n"
              f"Die mir vorliegende Septemberübersicht nennt {money(c['balance'])}. Bitte erläutern Sie die Zusammensetzung, soweit die Übersicht von unserem bisherigen Schriftwechsel abweicht. Ich bestätige mit dieser Nachricht weder die Berechtigung sämtlicher Nebenforderungen noch {contract_wording}. Sie erreichen mich schriftlich unter {c['address']} oder über diese E-Mail-Adresse.\n\nMit freundlichen Grüßen\n{correspondent}\n\n"
              f"Am 30.09.2026 schrieb Elin Voss:\nWir stellen die Unterlagen zu {stem} für das Projekt Lindenbogen zusammen. Bitte teilen Sie mit, welche offenen Punkte aus Ihrer Sicht mitgegeben werden müssen. Eine Zahlungsumstellung ist damit noch nicht verbunden.")
        eml(OUT/f'{stem}_04_Korrespondenz.eml',f'Re: {stem} Unterlagen und Überleitung',f'{correspondent} <{email}>','Elin Voss <elin.voss@ashcombe.example.invalid>',body,8,OUT/f'{stem}_03_Schreiben.pdf' if c['number'] in (7,8,10,13,17) else None)


def write_transaction(model):
    t=model['totals']
    docx(OUT/'00_Auftrag_Due_Diligence.docx','Prüfauftrag zum Projekt Lindenbogen',
         'An Rechtsanwältin Dr. Alva Steinacker. Wir möchten bis zum 16.10.2026 entscheiden, ob wir in verbindliche Verhandlungen über den Erfurter Darlehensbestand eintreten. Bitte beginnen Sie mit den 20 ausgewählten Einzelakten und dem Abgleich der Kaufpreisbasis.',[
      ('1 Auftrag und Rollen',[f"Ich, Selma Korngut, beauftrage Sie für {BUYER}, Pappelhof 9, 99084 Erfurt. Die Gesellschaft möchte 1.000 deutsche Verbraucherdarlehensforderungen von {BANK} erwerben. Auf Verkäuferseite koordiniert Elin Voss die Unterlagen, auf unserer Seite betreut Arvid Hummel die Zahlen. Bitte kennzeichnen Sie Verkäuferbehauptungen und eigene Befunde getrennt."]),
      ('2 Gegenstand des Erwerbs',["Unsere derzeitige Struktur ist ein Forderungskauf mit gesonderter Serviceüberleitung. Die Übernahme von Kundenverträgen als Ganzes, Einlagen, Filialmietvertrag, Personal oder Betriebsmitteln ist noch nicht vereinbart. Die Umgangssprache ‚Filialkauf‘ in den bisherigen Mails darf den Prüfungsgegenstand nicht erweitern. Wenn eine dieser Positionen für die geplante Abwicklung erforderlich ist, benötigen wir eine getrennte Entscheidung und einen passenden Nachweis."]),
      ('3 Prüfung und Ergebnisse',["Bitte erstellen Sie eine Management-Zusammenfassung, eine belegbezogene Feststellungsliste und konkrete Änderungen des Kaufvertragsentwurfs. Die Prüfung soll Bestand, Übertragbarkeit, Beweisbarkeit, Einwendungen, Durchsetzbarkeit, Servicing, Datenzugang, Kaufpreis und Aufsichtsvoraussetzungen verbinden. Einen Buchungssaldo wollen wir nicht ohne Weiteres als einziehbare Forderung bewerten. Schlagen Sie Nachforderungen und konkrete Vollzugsbedingungen vor.","Der Datenraum enthält 20 bewusst problemorientiert ausgewählte Verträge. Die übrigen 980 Datensätze sind lediglich Verkäuferexporte. Bitte leiten Sie daraus weder eine statistische Fehlerquote noch einen ungeprüften Befund für den Gesamtbestand ab. Wenn ein Risiko nicht nur den Einzelfall betrifft, benötigen wir eine begründete Erweiterung des Prüfungsumfangs."]),
      ('4 Zahlen und Vollzug',[f"Der Verkäuferexport weist eine Hauptforderung von {money(t['principal'])} und einen Gesamtsaldo von {money(t['balance'])} aus. Die Hauptbuchüberleitung nennt einen gesonderten Zahlungseingang von 720 Euro. Die Preisfaktoren in der Excel stammen allein vom Verkäufer. Für Einwendungen, Kosten und Sicherheiten sind zunächst keine bestätigten Erwerberwerte eingetragen.","Vor einer Freigabe müssen unsere Erlaubnis- und Servicingstruktur, die rechtswirksame Forderungsübertragung, die Datenübermittlung und die Kundenkommunikation geklärt sein. Ashcombe ist eine englische Gesellschaft. Bitte lassen Sie sich den tatsächlichen Status ihrer deutschen Tätigkeit und die für unsere Erwerbs- und Einziehungsstruktur erforderlichen Nachweise vorlegen. Wir unterstellen keinen fortbestehenden europäischen Pass."]),
      ('5 Kommunikation',["Fragen an den Verkäufer werden als Entwurf an mich gegeben. Das Team darf intern Unterlagen auswerten und Entwürfe erstellen; es darf keine Kunden anschreiben, Forderungen anerkennen, Kontoangaben umstellen oder Zahlungen anweisen. Bitte nennen Sie am Ende jeder Arbeitssitzung die nächste konkrete Entscheidung und die Unterlagen, die hierfür noch fehlen."])
    ],f'{BUYER}\nSelma Korngut | 10.10.2026')
    docx(OUT/'01_Letter_of_Intent_Entwurf.docx','Absichtserklärung zum Erwerb von Darlehensforderungen',
         f'{BANK} und {BUYER} halten den folgenden Verhandlungsstand vom 08.10.2026 fest. Dieser Entwurf beschreibt die beabsichtigte Transaktion; ein verbindlicher Forderungskauf wird erst durch einen gesondert freigegebenen Vertrag begründet.',[
      ('1 Transaktionsgegenstand',["Gegenstand der Verhandlungen sind die in der Verkäuferliste KREDIT-ERF-20260930 bezeichneten 1.000 Verbraucherdarlehensforderungen. Eine Übernahme der Bankgesellschaft, sämtlicher Kundenverträge oder des Filialbetriebs wird hierdurch nicht vereinbart. Parteienbezeichnungen und Vertretungsnachweise sind vor Unterzeichnung anhand vollständiger Register- und Vollmachtsunterlagen zu bestätigen."]),
      ('2 Preisgrundlage',[f"Der indikative Verkäuferansatz beträgt {money(t['price'])}. Er ergibt sich aus den Kapitalständen und den in der Arbeitsmappe offengelegten Faktoren; auf Zinsen und Kosten wird in diesem Ansatz kein zusätzlicher Kaufpreis gerechnet. Dieser Ansatz steht unter dem Vorbehalt der rechtlichen und wirtschaftlichen Prüfung. Durch bloße Einsicht in die Datei billigt die Käuferin weder den Datenbestand noch die Bewertung einzelner Forderungen."]),
      ('3 Prüfungsphase',["Die Verkäuferin stellt Vertragsfassungen, Auszahlungsbelege, Zahlungshistorien, Titel, Sicherheitenunterlagen, Einwendungen und Informationen zur regulatorischen und datenschutzrechtlichen Abwicklung zur Verfügung. Die Käuferin kann weitere Einzelakten anfordern. Die 20 bereitgestellten Detailakten werden nicht als repräsentative Zufallsstichprobe bezeichnet. Zugriff erhalten nur namentlich freigegebene Personen mit dokumentiertem Zweck."]),
      ('4 Geplanter Vollzug',["Als Arbeitsziel wird der 30.11.2026 genannt. Der Termin ist kein vorbehaltlos zugesagter Vollzugstag. Er setzt die abgestimmte Übertragungsstruktur, sämtliche erforderlichen Genehmigungen oder Erlaubnisnachweise, die abschließende Kaufpreismechanik, eine getestete Serviceüberleitung und eine freigegebene Kundeninformation voraus. Keine Partei darf durch diese Absichtserklärung fremde Rechte oder zwingende gesetzliche Voraussetzungen ersetzen."]),
      ('5 Vertraulichkeit und Umgang mit Daten',["Die Parteien verwenden bereitgestellte Unterlagen nur zur Prüfung und Vorbereitung dieser Transaktion. Personenbezogene Daten werden zunächst minimiert und nur soweit benötigt zugänglich gemacht. Jede Partei prüft ihre Rechtsgrundlage, Zugriffsberechtigungen und gegebenenfalls einschlägige Drittlandanforderungen. Eine Vertraulichkeitsvereinbarung allein ersetzt diese Prüfung nicht. Unberechtigter Zugriff ist der anderen Partei unverzüglich mit den bekannten Tatsachen mitzuteilen."]),
      ('6 Kosten und Verbindlichkeit',["Jede Partei trägt ihre eigenen Prüfungskosten. Die Regelungen über Vertraulichkeit, zweckgebundene Nutzung und Kosten sollen mit Unterzeichnung verbindlich sein; die übrigen Bestimmungen sind eine unverbindliche Verhandlungsgrundlage. Eine Exklusivität wird nicht vereinbart. Eine Partei kann die Verhandlungen beenden, muss dann aber die zweckgebundenen Unterlagen nach abgestimmten Aufbewahrungspflichten behandeln."]),
      ('7 Unterzeichnung',["Die Unterzeichnung dieses Entwurfs ist noch nicht freigegeben. Auf Verkäuferseite soll Beatrice Ashdown nach Vorlage ihrer Vertretungsbefugnis zeichnen; auf Käuferseite sollen Selma Korngut und Arvid Hummel gemeinsam zeichnen. Die interne Nennung ersetzt weder eine Vollmacht noch einen registerrechtlichen Nachweis."])
    ])
    spa=[
      ('1 Parteien und Vertragsstand',[f"Verkäuferin ist Ashcombe Personal Bank plc, handelnd durch ihre Niederlassung Erfurt, {ADDR}. Käuferin ist {BUYER}, Pappelhof 9, 99084 Erfurt. Die Verkäuferin erklärt, dass die Niederlassung keine eigene Rechtspersönlichkeit besitzt. Vertretungsnachweise und Registerauszüge sind vor Unterzeichnung in den Abschlussordner aufzunehmen. Dieser Text ist der Verhandlungsentwurf der Verkäuferseite vom 07.10.2026 und noch nicht angenommen."]),
      ('2 Kaufgegenstand und Abgrenzung',["Die Verkäuferin verkauft die in Anlage 1 eindeutig bezeichneten Forderungen gegen deutsche Verbraucher aus 1.000 Darlehensverträgen einschließlich der dort bezeichneten Nebenrechte. Anlage 1 ist die zum Vollzug von beiden Parteien abgezeichnete Endfassung des Bestandsverzeichnisses. Der derzeitige Mittagsauszug vom 30.09.2026 ist lediglich die Ausgangsfassung. Nicht enthalten sind Kundeneinlagen, sonstige Verbindlichkeiten, die Bankgesellschaft, Mietverträge, Arbeitsverhältnisse oder eine Erlaubnis zum Betrieb von Bankgeschäften.","Die Übernahme ganzer Kundenverträge wird nicht vereinbart. Soweit Rechte oder Pflichten ohne Mitwirkung Dritter nicht wie vorgesehen übergehen können, werden diese Positionen bis zur Klärung getrennt geführt. Eine Servicevereinbarung ersetzt weder die Forderungsübertragung noch eine erforderliche Zustimmung."]),
      ('3 Wirtschaftlicher Stichtag',["Wirtschaftlicher Stichtag soll der 30.09.2026, 24.00 Uhr, sein. Die Parteien gleichen alle Zahlungen zwischen dem Exportzeitpunkt um 12.00 Uhr und diesem Stichtag sowie die nachträglich wertgestellten Buchungen gesondert ab. Der Eingang von 720 Euro zu APB-0017 wird in einer Stichtagsbrücke offengelegt. Eine Zahlung darf weder zweimal preiswirksam abgezogen noch gleichzeitig beim Verkäufer und Käufer als einziehbare Forderung stehen bleiben."]),
      ('4 Kaufpreis und Überleitung',[f"Der vorläufige Kaufpreis beträgt {money(t['price'])}. Grundlage sind die Kapitalstände und Verkäuferfaktoren der Arbeitsmappe. Zinsen und Kosten sind im vorläufigen Preis nicht gesondert vergütet. Vor Unterzeichnung sind die zulässigen Ausschlüsse, Korrekturen und Einzelfallabschläge in Anlage 2 zu vereinbaren. Ein lediglich technisch vorhandener Buchungssaldo wird nicht durch diese Preisklausel rechtlich bestätigt.","Der endgültige Kaufpreis ergibt sich aus der gemeinsam festgestellten Endliste und den in Anlage 2 vereinbarten Korrekturen. Über streitige Positionen wird eine gesonderte Liste geführt. Bis zur Einigung werden diese Positionen nicht übertragen und nicht bezahlt. Die Parteien benennen für jede Position den Betrag, den Grund der Abweichung und die noch benötigten Belege."]),
      ('5 Abtretung und Übergabe',["Die Abtretung wird am Vollzugstag durch eine gesonderte, auf die Endliste Bezug nehmende Erklärung vorgenommen. Die Verkäuferin übergibt die vorhandenen Vertragsurkunden, Auszahlungsnachweise, Buchungshistorien, Einwendungen, Titel und Sicherheitenunterlagen. Fehlende Originale werden einzeln bezeichnet; ein allgemeiner Hinweis auf den Datenraum genügt nicht. Die Verkäuferin informiert über frühere Abtretungen, Verpfändungen und vertragliche Beschränkungen.","Die Parteien prüfen für jede nicht ohne Weiteres übertragbare Sicherheit oder Urkunde die erforderliche Mitwirkung und Form. Die Kennzeichnung ‚besichert‘ im System ist keine Bestätigung einer wirksam bestellten und überleitbaren Sicherheit. Einwendungen und Einreden der Schuldner werden durch den Verkauf nicht aufgehoben."]),
      ('6 Vollzugsbedingungen',["Der Vollzug setzt eine dokumentierte Entscheidung über die aufsichtsrechtlich zulässige Erwerbs- und Servicingstruktur, vollständige Vertretungsnachweise, die bereinigte Endliste und die freigegebene Zahlungsüberleitung voraus. Die Parteien dürfen nur solche Zustimmungen oder Bedingungen abbedingen, über die sie rechtlich verfügen können. Anforderungen von Behörden oder Rechte Dritter sind nicht durch eine interne Freigabe ersetzbar.","Die Verkäuferin legt den Nachweis ihres tatsächlichen Erlaubnis- und Niederlassungsstatus vor. Die Käuferin legt die für ihre konkrete Tätigkeit erforderlichen Nachweise vor. Ein Verweis auf eine frühere Tätigkeit innerhalb des europäischen Binnenmarkts ersetzt diese Nachweise nicht. Bis zur Erfüllung darf kein massenweiser Gläubigerwechsel an Kunden kommuniziert werden."]),
      ('7 Erklärungen der Verkäuferin',["Die Verkäuferin erklärt nach ihrer derzeitigen Kenntnis, Inhaberin der angebotenen Forderungen zu sein und die in der Endliste ausgewiesenen Zahlungseingänge vollständig offengelegt zu haben. Sie legt sämtliche ihr bekannten Einwendungen, Stundungen, Vergleiche, Insolvenzereignisse, Restschuldbefreiungen und Beschwerden offen. Jede Erklärung ist mit ihrem Wissensstand und den offengelegten Ausnahmen abzugleichen.","Der Umfang einer verschuldensunabhängigen Einstandspflicht, etwaige Wissenszurechnung sowie Rechtsfolgen falscher Angaben sind noch nicht abschließend verhandelt. Die Verkäuferin schlägt eine Haftungsobergrenze in Höhe von 15 Prozent des Kaufpreises und eine Anzeigefrist von 18 Monaten ab Vollzug vor. Die Käuferin hat diesen Vorschlag nicht akzeptiert; insbesondere Forderungsinhaberschaft, doppelte Verkäufe, vorsätzliche Täuschung und Datenschutzvorfälle sind getrennt zu behandeln."]),
      ('8 Zahlungseingänge und Fehlleitungen',["Nach Vollzug eingehende Zahlungen auf übertragene Forderungen werden täglich mit Vertragsnummer, Wertstellung und Zuordnung dokumentiert und nach dem vereinbarten Abrechnungsplan weitergeleitet. Eine unklare Zahlung bleibt bis zur Klärung sichtbar auf einer Ausnahmeliste. Rücklastschriften und Erstattungen werden derselben Forderung zugeordnet. Keine Partei darf einen belegten Eingang allein wegen einer technischen Verzögerung nochmals vom Kunden verlangen."]),
      ('9 Serviceüberleitung',["Die Verkäuferin soll für höchstens drei Monate nach Vollzug einfache Bestandsauskünfte und die technische Verbuchung als Übergangsleistung erbringen. Umfang, Weisungsbefugnisse, Entgelt, Datenschutzrollen, Kontrollrechte und etwaige Erlaubnisanforderungen werden vor Vollzug in einer gesonderten Servicevereinbarung festgelegt. Ohne deren Freigabe startet der Übergangsservice nicht. Kündigungen, Vergleiche, neue Vollstreckungsaufträge und Änderungen von Zahlungsdaten bedürfen einer namentlich dokumentierten Freigabe der zuständigen menschlichen Stelle."]),
      ('10 Kundenkommunikation',["Die Parteien stimmen eine verständliche Mitteilung ab, die Forderungsinhaber, Servicer, unveränderte Konditionen, Ansprechpartner und den überprüfbaren Zahlungsweg auseinanderhält. Betroffene mit Insolvenzverfahren, streitiger Identität, unklarer Erbfolge oder besonderen Vereinbarungen erhalten keine unbesehene Standardnachricht. Widersprüchliche Anschriften und Zustellrückläufer sind vor Versand zu klären. Die Mitteilung enthält keine Aufforderung zur Preisgabe von PIN oder Zugangsdaten."]),
      ('11 Daten und Informationssicherheit',["Die Parteien dokumentieren Zweck, Umfang, Rechtsgrundlage und Empfänger jeder Datenweitergabe. Zugriffe werden rollenbezogen beschränkt und protokolliert. Ausweis-, Gesundheits- und Einkommensunterlagen werden nicht allein deshalb vollständig kopiert, weil sie im Altbestand vorhanden sind. Der mögliche Portalvorfall bei APB-0015 wird gesondert untersucht; bekannte Risiken dürfen im Abschlussordner nicht durch einen pauschalen Vertraulichkeitshinweis ersetzt werden.","Für den Übergangsservice in Manchester werden tatsächliche Zugriffswege, Datenschutzrollen und erforderliche Übermittlungsinstrumente vor Freigabe geprüft. Eine bloße Vertragsklausel behauptet keine Angemessenheitsentscheidung oder gesetzliche Zulässigkeit. Aufbewahrung, Rückgabe und Löschung werden nach Zweck und geltenden Pflichten abgestimmt; Beweismittel in offenen Streitigkeiten dürfen nicht unbesehen gelöscht werden."]),
      ('12 Besondere Bestandsfälle',["APB-0001, APB-0002, APB-0008, APB-0014 und APB-0019 werden bis zur Entscheidung über Insolvenzstatus, Restschuldbefreiung, Abtretungsbeschränkung, Identität beziehungsweise Vergleich gesondert geführt. Diese Nennung ist keine abschließende Risikoliste. Die Käuferin kann aufgrund nachvollziehbarer Befunde weitere Positionen zur Einzelklärung benennen. Eine fehlende Dokumentation wird nicht durch ein unangekreuztes Datenfeld als negativer Befund bestätigt."]),
      ('13 Auseinandersetzung über Abweichungen',["Streitige Bestands- oder Kaufpreispositionen werden zunächst durch die benannten Projektverantwortlichen anhand der Originalbelege abgeglichen. Bleibt eine Differenz bestehen, wird sie mit beiden Positionen und der zugrunde liegenden Berechnung dokumentiert. Ein Sachverständiger kann nur für ausdrücklich vereinbarte Rechenfragen eingeschaltet werden; rechtliche Streitfragen werden nicht ohne gesonderte Vereinbarung seiner abschließenden Entscheidung unterstellt."]),
      ('14 Abschlussunterlagen und Unterzeichnung',["Zum Vollzug gehören Endliste, Kaufpreisüberleitung, Abtretungserklärung, Urkundeninventar, offene Ausnahmen, Servicevereinbarung und freigegebene Kundenmitteilungen. Die Parteien protokollieren Datum, verantwortliche Personen und die tatsächlich übergebenen Fassungen. Entwurfsdateien im Datenraum ersetzen keine Unterzeichnung. Beatrice Ashdown soll für die Verkäuferin zeichnen, Selma Korngut und Arvid Hummel sollen für die Käuferin gemeinsam zeichnen; ihre Befugnisse sind noch nachzuweisen."])
    ]
    docx(OUT/'02_Forderungskauf_und_Service_Entwurf.docx','Forderungskaufvertrag Projekt Lindenbogen',
         'Verkäuferentwurf vom 07.10.2026. Die Käuferseite soll den Text anhand der Due Diligence bearbeiten. Offene Verhandlungspunkte sind in vollständigen Klauseln beschrieben; sie sind keine bereits erteilten Freigaben.',spa)
    pdf(OUT/'03_Hauptbuch_und_Stichtagsbruecke.pdf','Hauptbuchüberleitung zum 30 September 2026',[
      f"Alma Wurzel, Buchhaltung Erfurt | 02.10.2026. Das Darlehensnebenbuch wurde am 30.09.2026 um 12.00 Uhr exportiert. Sein Gesamtbetrag beträgt {money(t['balance'])}. Darin enthalten sind Kapital {money(t['principal'])}, Zinsen {money(t['interest'])} und Kosten {money(t['costs'])}.",
      "Auf dem Zahlungssammelkonto 1690 ging am 30.09.2026 um 16.42 Uhr ein Betrag von 720 Euro ein. Der Verwendungszweck und der Zahler passen zu APB-0017. Die Zuordnung erfolgte technisch am 01.10.2026 mit Wertstellung 30.09.2026. Die Zahlung darf im Stichtagsabgleich nur einmal berücksichtigt werden.",
      f"Für den hier abgegrenzten Bestand ergibt sich damit eine technische Nettoposition von {money(model['general_ledger_net'])}. Dieser Abgleich betrifft den Buchungsbestand, nicht dessen Werthaltigkeit. Einzelwertberichtigungen und der regulatorische Abschluss sind in dieser Datei nicht enthalten. Eine vollständige Bankbilanz wird hiermit nicht vorgelegt.",
      "Die Abgrenzung enthält keine Kundeneinlagen. Die Hauptbuchkonten wurden ausschließlich für die angebotenen 1.000 Forderungen gefiltert. Die Daten stammen aus einem Verkäuferexport. Der Erwerber soll den Filter, den Zugriff auf Originalbelege und die Zuordnung der Zahlung noch nachvollziehen."
    ],BANK,(['Position','Betrag'],[['Nebenbuch Export 12.00 Uhr',money(t['balance'])],['Zahlung auf Sammelkonto',money(-72000)],['Nettoposition nach Überleitung',money(model['general_ledger_net'])]],[115,48]))
    pdf(OUT/'04_Datenraum_und_Urkundeninventar.pdf','Datenraumstand und Urkundeninventar',[
      "Elin Voss an das Projektteam | 09.10.2026. Die Ordnernummern APB-0001 bis APB-0020 sind bewusst ausgewählte Detailfälle. Sie decken unterschiedliche Bearbeitungsverläufe ab. Die Auswahl erfolgte durch den Verkäufer und ist weder zufällig noch repräsentativ. Für APB-0021 bis APB-1000 stehen derzeit nur die knappen Exportdaten zur Verfügung.",
      "Jede Detailakte enthält Vertragsabschrift, Kontojournal, ein wesentliches Schreiben und eine aktuelle E-Mail. Bei fünf E-Mails ist das genannte Schreiben zusätzlich als PDF-Anhang eingebettet. Die Arbeitsmappe enthält die 1.000 Vertragsnummern sowie die zwölf Monatsbuchungen jedes Detailfalls. Frühere Buchungen sind als übernommene Anfangssalden ausgewiesen; die Originalbelege können nachgefordert werden.",
      "Nicht durchgehend vorhanden sind Signaturprotokolle, Aushändigungsnachweise, vollständige alte Informationsblätter, Originaltitel und Sicherheitenurkunden. APB-0007 enthält eine archivierte Vertragsabschrift, aber keinen belegten Zugang der Widerrufsinformation. APB-0012 enthält die Behauptung einer Sicherheit ohne vollständige Urkundenkette. APB-0020 wurde durch eine spätere unterschriebene Abschrift ergänzt.",
      "Im Bereich Unternehmensunterlagen sind Vertretung, tatsächlicher aufsichtsrechtlicher Status und die rechtliche Ausgestaltung des Übergangsservices noch nicht dokumentiert. Die Verkäuferseite verwendet intern den Begriff Filialverkauf; die aktuelle Transaktionsliste bezeichnet ausschließlich Forderungen. Personal, Mietvertrag und Einlagen wurden nicht als Kaufgegenstände freigegeben."
    ])
    # Lebensnahe Korrespondenz zu Deal, Zahlen, Regulatorik, Zeitdruck und Kundeninformation.
    mails=[
      ('05_Verkaeufer_Zeitplan.eml','Projekt Lindenbogen: Unterlagen und Zeitplan','Beatrice Ashdown <beatrice.ashdown@ashcombe.example.invalid>','Selma Korngut <selma.korngut@ilmgrund.example.invalid>',"Sehr geehrte Frau Korngut,\n\nwir möchten die Erfurter Aktivität zum Jahresende verkleinern und bitten um ein indikatives Angebot bis zum 16.10. Die beigefügten 20 Akten stammen aus der Bearbeitung des Teams; ich kann nicht bestätigen, dass die Häufigkeit der Probleme dem Gesamtbestand entspricht. Die 980 weiteren Datensätze sind noch nicht einzeln kopiert. Wenn Ihre Berater weitere Unterlagen benötigen, benennen Sie bitte die Nummern und den Grund, damit wir die Zugriffe begrenzen können.\n\nUnser CFO spricht in den Präsentationen von einem Filialverkauf. Ich verstehe darunter im ersten Schritt nur die Forderungen. Mietvertrag, Beschäftigte und Einlagen sind bislang nicht Gegenstand einer freigegebenen Vereinbarung. Bitte lassen Sie uns diese Begriffe vor einer Kundenmitteilung bereinigen. Die beiden Ansprechpartner für Originalurkunden sind bis Mittwoch teilweise im Urlaub. Das sollte uns nicht zu einer pauschalen Bestätigung unbekannter Unterlagen verleiten.\n\nMit freundlichen Grüßen\nBeatrice Ashdown"),
      ('06_Kaeufer_QA.eml','Datenraumfragen 1 bis 7 und nächste Freigabe','Selma Korngut <selma.korngut@ilmgrund.example.invalid>','Elin Voss <elin.voss@ashcombe.example.invalid>',"Sehr geehrte Frau Voss,\n\nwir benötigen erstens die Herleitung des vollständigen Bestands von 1.000 Forderungen und zweitens die unveränderte Mittagsdatei samt Stichtagsbrücke. Drittens bitten wir um die Originale beziehungsweise nachvollziehbaren Kopien der Titel und Sicherheiten. Viertens fehlen uns bei mehreren Altverträgen vollständige Aushändigungs- und Signaturprotokolle. Fünftens benötigen wir Ihre Zusammenstellung sämtlicher wirksamer Stundungen und Vergleiche. Sechstens bitten wir um die tatsächlichen Erlaubnis- und Vertretungsnachweise für die englische Bank und ihre deutsche Tätigkeit. Siebtens brauchen wir die Beschreibung des vorgesehenen Services in Manchester einschließlich der Datenzugriffe.\n\nBitte beantworten Sie die Fragen jeweils mit einem konkreten Dokument oder einer ausdrücklichen Fehlanzeige. Die Aussage ‚steht im System‘ reicht für die Prüfung nicht. Wir möchten anschließend über Ausschlüsse, Kaufpreis und Vollzugsbedingungen verhandeln. Eine Zahlung oder Kundenanschriftänderung ist derzeit nicht freigegeben.\n\nMit freundlichen Grüßen\nSelma Korngut"),
      ('07_Buchhaltung_Nachmittag.eml','720 Euro nach dem Export sind kein zusätzlicher Ertrag','Alma Wurzel <alma.wurzel@ashcombe.example.invalid>','Arvid Hummel <arvid.hummel@ilmgrund.example.invalid>',"Sehr geehrter Herr Hummel,\n\nanbei die Überleitung. Bitte reduzieren Sie nicht sowohl das Kapital im ursprünglichen Export als auch nochmals den Kaufpreis um dieselben 720 Euro. Der Zahlungseingang steht zunächst auf dem Sammelkonto; die konkrete Aufteilung auf Zins und Kapital ist noch abzugleichen. Ich habe deshalb eine Stichtagsbrücke auf Gesamtsaldoebene erstellt und die Ursprungstabelle unverändert gelassen. Die Preisfaktoren beziehen sich dagegen auf Kapital. Wie sich die Zuordnung auf den endgültigen Preis auswirkt, muss die Endliste zeigen.\n\nUnsere Tabelle enthält keine Einzelwertberichtigungen. Sie ist damit keine Unternehmensbewertung und keine Bankbilanz. Die Kostenpositionen wurden aus dem Nebenbuch übernommen, auch wenn Kunden sie bestreiten. Bitte nennen Sie sie im Bericht nicht bereits unstreitig geschuldet. Für die Dokumentation können wir Ihnen den Kontenauszug des Sammelkontos mit geschwärzten Fremdpositionen nachreichen.\n\nMit freundlichen Grüßen\nAlma Wurzel"),
      ('08_Service_Manchester.eml','Serviceüberleitung: Berechtigungen und Postfach','Ruth Caldwell <ruth.caldwell@ashcombe.example.invalid>','Selma Korngut <selma.korngut@ilmgrund.example.invalid>',"Sehr geehrte Frau Korngut,\n\ndas Team in Manchester könnte für drei Monate Buchungsdateien erzeugen und Kundenanfragen vorsortieren. Wir benötigen dafür nicht automatisch den gesamten bisherigen Dokumentenbestand. Bitte teilen Sie uns die festgelegten Aufgaben und Ihre Freigabeverantwortlichen mit. Rechtliche Entscheidungen über Kündigung, Insolvenz oder Vergleich sollen nicht durch eine automatische Statusregel ausgelöst werden.\n\nUnsere technische Skizze nennt eine gemeinsame Ablage. Ob beide Seiten dafür dieselbe datenschutzrechtliche Rolle haben, ist nicht entschieden. Ebenso offen sind der genaue Empfängerkreis und die Zugriffe aus dem Vereinigten Königreich. Die IT kann eine Verbindung einrichten, aber damit ist keine Rechtsgrundlage bestätigt. Zum Portalvorfall APB-0015 existiert bisher nur das Kundenticket; das vollständige Zugriffsprotokoll ist angefordert. Bitte planen Sie den Servicebeginn erst nach einer gemeinsamen fachlichen und rechtlichen Abnahme.\n\nMit freundlichen Grüßen\nRuth Caldwell"),
      ('09_Erlaubnis_und_Vertretung.eml','Erlaubnisstatus und Zeichnung sind noch offen','Dr. Jonna Hagedorn <jonna.hagedorn@ilmgrund.example.invalid>','Selma Korngut <selma.korngut@ilmgrund.example.invalid>',"Sehr geehrte Frau Korngut,\n\ndie derzeitigen Projektunterlagen enthalten weder eine belastbare Erlaubnisakte der Verkäuferin noch unsere abschließende Einordnung der geplanten Tätigkeit. Bitte unterscheiden Sie das Halten erworbener Forderungen, das konkrete Servicing und etwaige darüber hinausgehende Bankgeschäfte. Die Gründung einer deutschen GmbH allein beantwortet diese Fragen nicht. Ebenso wenig genügt die Aussage, Ashcombe sei seit vielen Jahren in Deutschland tätig.\n\nIch habe um die aktuellen Nachweise sowie um die Beschreibung der notleidenden Teilbestände gebeten. Die Unterlagen sollen vor einer Entscheidung vollständig vorliegen. Beatrice Ashdown hat als Ansprechpartnerin geschrieben, aber ihre Vertretungsbefugnis ist noch nicht nachgewiesen. Bitte versenden Sie den Kaufvertrag nicht als angeblich bereits unterschriftsreif. Die notwendige Prüfung kann Auswirkungen auf Erwerberstruktur, Servicevertrag und Vollzugsbedingungen haben.\n\nMit freundlichen Grüßen\nDr. Jonna Hagedorn"),
      ('10_Kundeninformation_Entwurf.eml','Entwurf der Kundenmitteilung bitte noch nicht versenden','Elin Voss <elin.voss@ashcombe.example.invalid>','Selma Korngut <selma.korngut@ilmgrund.example.invalid>',"Sehr geehrte Frau Korngut,\n\nim ersten Entwurf stand: ‚Ihre Filiale gehört jetzt Ilmgrund, überweisen Sie ab sofort auf das neue Konto.‘ Diesen Satz habe ich vorerst gestrichen. Der Vertrag ist nicht unterzeichnet, und für die problematischen Einzelfälle stehen weder Empfänger noch Zahlungsweg fest. Bitte geben Sie einen Text frei, der Gläubiger und Servicer auseinanderhält und nur tatsächlich vollzogene Änderungen mitteilt.\n\nFür APB-0001 soll die Insolvenzverwaltung angeschrieben werden; bei APB-0010 ist die Erbfolge ungeklärt. APB-0004 hat einen Adresskonflikt, APB-0014 bestreitet die Identität. Diese vier Gruppen dürfen nicht einfach in den Standardversand laufen. Wir brauchen eine Empfängerliste, einen dokumentierten Abgleich und eine namentliche Freigabe. Die Kunden sollen keine PIN eingeben und keine unbekannte Zahlungsseite aufrufen müssen. Den Versand beginnen wir erst nach Ihrer Rückmeldung und dem dokumentierten Vollzug.\n\nMit freundlichen Grüßen\nElin Voss")
    ]
    for fn,subject,sender,to,body in mails:
        eml(OUT/fn,subject,sender,to,body,9,OUT/'03_Hauptbuch_und_Stichtagsbruecke.pdf' if fn.startswith('07') else None)

    a=model['branch_accounts']
    pdf(OUT/'12_Filialergebnis_und_Buchwerte.pdf','Erfurter Ergebnisrechnung und Bestandsbuchwerte',[
      'Alma Wurzel an Arvid Hummel | 09.10.2026. Diese interne Zusammenstellung zeigt die dem Erfurter Kreditbestand zugeordneten Ergebnispositionen. Sie ist weder ein testierter Jahresabschluss noch eine vollständige Bilanz der englischen Bank. Das Jahr 2025 umfasst zwölf Monate, die Spalte 2026 nur Januar bis September; ein unmittelbarer Jahresvergleich wäre irreführend.',
      f"Die technische Nettoposition nach Berücksichtigung des Zahlungssammelkontos beträgt {money(model['general_ledger_net'])}. Davon wurden {money(a['impairment_individual'])} einzeln zugeordnete und {money(a['impairment_portfolio'])} portfoliobezogene Wertberichtigungen abgezogen. Daraus ergibt sich ein interner Buchwert von {money(a['book_value'])}. Die vollständige Zuordnung der Einzelwertberichtigungen auf Vertragsnummern ist noch nicht aus dem Altsystem exportiert. Bitte setzen Sie diesen Buchwert nicht mit dem verhandelten Kaufpreis oder einer bestätigten wirtschaftlichen Bewertung gleich.",
      'In den Personalkosten sind fünf Beschäftigte enthalten. Ihr möglicher Übergang wurde bisher nicht als Teil des Forderungskaufs vereinbart. Der Mietvertrag läuft nach der Übersicht der Filialleitung bis 31.12.2028; eine Übertragung oder vorzeitige Beendigung ist nicht bestätigt. IT-Kosten und zentrale Umlagen entfallen nach einem Verkauf nicht automatisch. Der Serviceentwurf muss deshalb zwischen entfallenden Kosten, fortlaufenden Restkosten und neuem Übergangsaufwand unterscheiden.',
      'Die Risikovorsorge der Ergebnisrechnung ist eine Periodenbewegung. Sie darf nicht nochmals als zusätzlicher Stichtagsabzug neben dem bereits verminderten Buchwert berücksichtigt werden. Die nachfolgende Übersicht erlaubt lediglich eine erste Nachforderungsliste für das finanzielle Prüfungsteam.'
    ],table=(['Position','2025 vollständig','2026 Januar bis September'],[[label,money(x),money(y)] for label,x,y in zip(a['categories'],a['income_2025'],a['income_2026_ytd'])]+[['Ergebnis vor weiteren Steuern und Konzernpositionen',money(sum(a['income_2025'])),money(sum(a['income_2026_ytd']))]],[73,43,47]))
    c=model['cases'][0]
    p=c['journal'][7]['principal_end'];ins_interest=c['journal'][7]['interest_end']
    pdf(OUT/'APB-0001_05_Tabellenauszug_Verwalter.pdf','Auszug aus der Forderungsliste Schlehenried',[
      'Dr. Malte Hohenacker, Insolvenzverwaltung | 24.09.2026. Auszug aus der internen Arbeitsliste nach Anmeldung vom 26.06.2026. Schuldnerin: Kunigunde Schlehenried. Eröffnung: 15.06.2026. Die Liste gibt den Bearbeitungsstand wieder und ersetzt keinen beglaubigten gerichtlichen Tabellenauszug.',
      f"Gläubigerin ist Ashcombe Personal Bank plc, Niederlassung Erfurt. Vertragsnummer APB-0001. In der Anmeldung sind {money(p)} Kapital und {money(ins_interest)} bis 31.05.2026 gebuchte unbezahlte Zinsen bezeichnet. Die Verkäuferin hat zusätzlich 36 Euro Mahnkosten angemeldet. Für den Zeitraum bis zur Eröffnung am 15.06.2026 wurde keine gesonderte Zinsberechnung beigefügt; diese Position bleibt nachzufordern.",
      'Die Schuldnerin bestreitet die Mahnkosten. In der vorliegenden Liste ist die Kapitalforderung als nicht bestritten vermerkt; eine gerichtliche Feststellung oder vollstreckbare Ausfertigung ist dieser internen Datei nicht beigefügt. Eine Forderung aus vorsätzlich begangener unerlaubter Handlung wurde nicht angemeldet. Gesonderte Sicherheitenurkunden liegen uns nicht vor.',
      'Ein künftiger Forderungserwerber soll nicht allein durch Übersendung einer Excel-Zeile in die Gläubigerkorrespondenz aufgenommen werden. Wir benötigen die dokumentierte Rechtsnachfolge und eine klare Bezeichnung der betroffenen Forderung. Bitte übermitteln Sie nach Eröffnung gebuchte Zinsen getrennt von der vorliegenden Anmeldung.'
    ],table=(['Position','Betrag','Vermerk'],[['Kapital per 31.05.2026',money(p),'Arbeitsstand nicht bestritten'],['Zinsen bis 31.05.2026',money(ins_interest),'Nachtragsberechnung angefordert'],['Mahnkosten',money(3600),'Von Schuldnerin bestritten']],[70,36,57]))
    pdf(OUT/'APB-0005_05_Titel_und_Zahlungsnachweis.pdf','Titeldaten und Abrechnung des Vollstreckungsmandats',[
      'Lorenz Tannreuther an Ashcombe | 30.09.2026. Übertragung aus der Kanzleiakte, Vorgang APB-0005. Schuldnerin: Ottilie Haberkamm, Kirschanger 4, 99091 Erfurt. Der Titel wird in unserem Urkundenregister als Vollstreckungsbescheid vom 28.08.2026 geführt; der eingescannte Teil im Bankdatenraum enthält die nachfolgende Forderungsaufstellung, jedoch nicht den vollständigen Zustellvermerk.',
      'Der Titel betrifft eine bezifferte Teilforderung von 1.500 Euro aus den Tilgungsanteilen Mai bis Juli 2026 nebst darin aufgeführten Verfahrenskosten von 186 Euro. Eine Vollstreckung des gesamten laufenden Kapitalsaldos wurde damit nicht tituliert. Der Originaltitel liegt in unserer Handakte. Vor einer weiteren Maßnahme gleichen wir Hauptforderung, titulierte Kosten, tatsächlich entstandene Vollstreckungskosten und bereits vereinnahmte Zahlungen ab.',
      'Am 16.09.2026 gingen 180 Euro von Frau Haberkamm auf unserem Fremdgeldkonto ein, Referenz VH-0916-APB0005. Am 18.09.2026 wurden genau 180 Euro an die Bank weitergeleitet, Referenz FK-0918-APB0005. Die Kundin hatte eine Tilgungsbestimmung zugunsten der Kapitalforderung abgegeben; die Bank bestätigte diese konkrete Zuordnung am 21.09.2026. Die Zahlung erscheint deshalb im Septemberjournal auf Kapital. Sie ist kein zweiter zusätzlicher Zahlungseingang.',
      'Die im Bankjournal weiter geführten 120 Euro sind keine von unserer Kanzlei abgerechneten Vollstreckungskosten. Bitte trennen Sie diese interne Bankposition von den titulierten 186 Euro und dem im Titel ausgewiesenen Anspruch. Die vollständige Titelkopie einschließlich Zustellung und etwaiger Klausel ist vor Übergabe an den Erwerber gesondert zu vervollständigen.'
    ],table=(['Datum','Beleg','Betrag'],[['16.09.2026','Kundeneingang VH-0916-APB0005',money(18000)],['18.09.2026','Weiterleitung FK-0918-APB0005',money(18000)],['21.09.2026','Bestätigte Zuordnung Kapital',money(18000)]],[35,86,42]))


def write_readme(model):
    OUT.joinpath('README.md').write_text(f'''# Due Diligence einer Erfurter Bankfiliale

Die Akte begleitet den beabsichtigten Erwerb von 1.000 Verbraucherdarlehensforderungen einer fiktiven englischen Bank. Sie eignet sich für Käuferprüfung, Verkäuferaufbereitung, Kaufpreisabgleich und die Bearbeitung des Forderungskaufvertrags. Ausgangspunkt ist ein Forderungskauf mit gesonderter Serviceüberleitung; Einlagen, vollständige Kundenverträge und Betriebsteile sind keine stillschweigend mitverkauften Gegenstände.

## 1 Benutzung

Beginnen Sie mit `00_Auftrag_Due_Diligence.docx`, der Arbeitsmappe und dem Verkäuferentwurf. Bearbeitungsdatum ist der 10.10.2026, wirtschaftlicher Stichtag der 30.09.2026. Prüfen Sie aus Käufer- oder Verkäufersicht und verlangen Sie einen belegten Befund mit nächstem Arbeitsschritt. Eine fertige DD-Musterlösung ist nicht enthalten.

Die 20 Detailfälle wurden bewusst problemorientiert ausgewählt. Sie bilden keine repräsentative Stichprobe der 1.000 Forderungen. Für weitere 980 Nummern existieren nur knappe Verkäuferdaten; deren Einzelakten müssen nachgefordert werden. Alle Zahlen und Personennamen sind erfunden. Anschriften sind fiktiv; E-Mail-Adressen verwenden die nicht zustellbare Domain `example.invalid`. Es werden keine echten Bankkennzeichen, Kontoverbindungen oder Registereintragungen behauptet.

## 2 Unterlagen

- 20 ausformulierte Verbraucherdarlehensverträge in Word.
- 20 PDF-Kontojournale mit je zwölf Monatsbuchungen und ausdrücklich bezeichneten Anfangssalden.
- 20 konkrete Briefe in PDF und 20 E-Mails; fünf enthalten das zugehörige Schreiben als echten PDF-Anhang.
- Prüfauftrag, Absichtserklärung und bearbeitbarer Forderungskaufvertrag in Word.
- Hauptbuchüberleitung, Urkundeninventar und sechs Transaktions-E-Mails.
- Interne Filialergebnisrechnung mit Buchwertüberleitung sowie konkrete Insolvenz- und Vollstreckungsbelege.
- Excel mit 1.000 Forderungen, 240 Detailbuchungen, Stichtagsbrücke, Preisannahmen und Erläuterung der Belege.

Die PDF-Briefe verwenden die PDF-Standardschrift Times. Word ist auf Times New Roman 11 pt eingestellt; der verfügbare PDF-Renderer ersetzt die nicht installierte Schrift durch Liberation Serif. Die technische Darstellung eines Buchungssaldos bestätigt weder Fälligkeit noch Durchsetzbarkeit. Kosten, Zinsen, Insolvenzereignisse und kundenseitige Einwendungen bleiben sichtbar. Verkäuferfaktoren sind ausdrücklich keine Erwerberbewertung.

## 3 Downloads

<!-- BEGIN gesamt-pdf-section (autogen) -->

> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.
>
> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.

| Fassung | Download |
| --- | --- |
| Gesamt-PDF | [Gesamte Akte](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/due-diligence-v1.0.0/{SLUG}_gesamt.pdf) · [Repositoryfassung](gesamt-pdf/{SLUG}_gesamt.pdf) |
| Originaldateien | [Akten-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/due-diligence-v1.0.0/testakte-{SLUG}.zip) |
| Einzel-PDFs | [Einzel-PDF-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/due-diligence-v1.0.0/testakte-{SLUG}-einzelpdfs.zip) |

<!-- END gesamt-pdf-section (autogen) -->

## 4 Herkunft und Reproduktion

Die Unterlagen werden mit `scripts/build-dd-bankakte.py` und `scripts/build-dd-bankakte-excel.mjs` aus dem deterministischen Fallmodell reproduziert. Das Modell enthält Buchungsdaten, ausdrücklich getrennte Vorannahmen und die Bestandsüberleitung. Die Darstellung behauptet keine tatsächliche Banklizenz und ersetzt keine aufsichtsrechtliche Prüfung des Erwerbers oder Servicers.
''',encoding='utf-8')


def main():
    p=argparse.ArgumentParser();p.add_argument('--model-only',action='store_true');args=p.parse_args()
    model=build_model();MODEL.parent.mkdir(parents=True,exist_ok=True)
    MODEL.write_text(json.dumps(model,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    if not args.model_only:
        OUT.mkdir(parents=True,exist_ok=True);write_cases(model);write_transaction(model);write_readme(model)
    print(json.dumps({'cases':20,'portfolio':1000,'files':len(list(OUT.glob('*'))) if OUT.exists() else 0,'totals_cents':model['totals']},ensure_ascii=False))


if __name__=='__main__':main()
