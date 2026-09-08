#!/usr/bin/env python3
"""Prüft die 20 AML-Fachwege und drei getrennten Aktenfassungen."""

import importlib.util
import io
import json
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path
import re
import tempfile
import unittest
import zipfile

from docx import Document
from pypdf import PdfReader
from testakte_disclaimer import NOTICE_BYTES, pdf_content_errors
from testakte_einzelpdf_common import expected_arcnames

ROOT = Path(__file__).resolve().parent.parent
PLUGIN = ROOT / 'geldwaeschepraevention-aml-kyc'
CASES = ('aml-kanzlei-fremdgeld-anteilskauf-bonn',
         'aml-notariat-kaufpreis-drittzahlung-muenster',
         'aml-gueterhandel-teilzahlungen-maschinenhandel-ulm')


def builder(filename):
    spec = importlib.util.spec_from_file_location(filename.replace('-', '_'), ROOT / 'scripts' / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def skill(name):
    return (PLUGIN / 'skills' / name / 'SKILL.md').read_text(encoding='utf-8')


class GeldwaescheTests(unittest.TestCase):
    def test_twenty_specific_skills_and_contained_references(self):
        files = list((PLUGIN / 'skills').glob('*/SKILL.md'))
        self.assertEqual(len(files), 20)
        for path in files:
            text = path.read_text(encoding='utf-8')
            with self.subTest(skill=path.parent.name):
                self.assertIn('Times New Roman 11 pt', text)
                self.assertIn('## 6. Beispiele', text)
                self.assertIn('### 3.1.', text)
                self.assertNotIn(chr(167), text)
                desc = re.search(r'^description: (.+)$', text, re.M).group(1)
                self.assertNotRegex(desc, r'Pruef|Geldwaesche|fuer |Bussgeld')
                for target in re.findall(r'\]\(([^)]+)\)', text):
                    if '://' in target or target.startswith('#'):
                        continue
                    resolved = (path.parent / target.partition('#')[0]).resolve()
                    self.assertTrue(resolved.is_relative_to(PLUGIN.resolve()), target)
                    self.assertTrue(resolved.is_file(), target)

    def test_migration_targets_exist_without_wrappers(self):
        rows = [r for r in json.loads((ROOT / 'scripts/skill-selection-migrations.json').read_text()) if r['plugin'] == PLUGIN.name]
        self.assertEqual(len(rows), 41)
        self.assertEqual(len({r['old'] for r in rows}), 41)
        for row in rows:
            self.assertFalse((PLUGIN / 'skills' / row['old'] / 'SKILL.md').exists())
            self.assertTrue((PLUGIN / 'skills' / row['new'] / 'SKILL.md').exists())

    def test_present_and_future_law_are_separate(self):
        for kind in ('werkstatt', 'schnellstart'):
            text = (PLUGIN / f'{PLUGIN.name}-{kind}.md').read_text(encoding='utf-8')
            for term in ('2026', '2027', 'GwGMeldV', '16a', '43', '46', '2024/1624'):
                self.assertIn(term, text)
            self.assertNotIn('skills/', text)
            if kind == 'schnellstart':
                self.assertLess(len(text.encode('utf-8')), 7500)
            else:
                self.assertGreater(len(text.encode('utf-8')), 25000)
        self.assertIn('keine kleinere Barzahlung', skill('notariat-immobilienzahlung-pruefen'))
        self.assertIn('Fünf-Werktage', skill('notariat-immobilienzahlung-pruefen'))
        self.assertIn('seit 1. März 2026', skill('notariat-immobilienzahlung-pruefen'))
        self.assertIn('Der Skill sendet nichts selbst', skill('aml-verdachtsmeldung-fiu-leitfaden'))

    def test_eight_native_documents_with_complete_emails(self):
        for slug in CASES:
            case = ROOT / 'testakten' / slug
            sources = [p for p in case.iterdir() if p.is_file() and p.name[:2].isdigit()]
            self.assertEqual(len(sources), 8)
            self.assertEqual(sorted(p.suffix for p in sources), ['.docx'] + ['.eml'] * 4 + ['.txt'] * 3)
            for path in sources:
                if path.suffix == '.docx':
                    text = '\n'.join(p.text for p in Document(path).paragraphs)
                elif path.suffix == '.eml':
                    msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                    for header in ('From', 'To', 'Date', 'Subject', 'Message-ID', 'MIME-Version', 'Content-Type'):
                        self.assertTrue(msg[header], (path, header))
                    self.assertEqual(parsedate_to_datetime(msg['Date']).year, 2026)
                    text = msg.get_content()
                else:
                    text = path.read_text(encoding='utf-8')
                self.assertGreater(len(text), 500, path)
                self.assertNotRegex(text.lower(), r'musterlösung|lösung(smatrix|svorschlag)|arbeitsauftrag|formathinweis')
                self.assertNotIn(chr(167), text)

    def test_amounts_match_source_records(self):
        ulm = ROOT / 'testakten' / CASES[2]
        cash = (ulm / '02_kassenblatt.txt').read_text()
        for amount in ('4000', '3500', '3000', '10500'):
            self.assertIn(amount, cash)
        text = '\n'.join(p.text for p in Document(ulm / '01_auftragsbestaetigung.docx').paragraphs)
        for amount in ('25000', '4750', '29750'):
            self.assertIn(amount, text.replace(' ', '').replace('.', ''))
        self.assertEqual(4000 + 3500 + 3000 + 29750 - (25000 + 4750), 10500)
        muenster = ROOT / 'testakten' / CASES[1]
        text = '\n'.join(p.text for p in Document(muenster / '01_verkaeufererklaerung.docx').paragraphs)
        self.assertIn('486000', text.replace(' ', '').replace('.', ''))
        self.assertEqual(350000 + 128000 + 8000, 486000)

    def test_combined_pdf_contains_all_sources(self):
        for slug in CASES:
            case = ROOT / 'testakten' / slug
            pdf = case / 'gesamt-pdf' / f'{slug}_gesamt.pdf'
            self.assertEqual(pdf_content_errors(pdf.read_bytes()), [])
            text = '\n'.join(p.extract_text() or '' for p in PdfReader(pdf).pages)
            for source in case.iterdir():
                if source.is_file() and source.name[:2].isdigit():
                    self.assertIn(source.name, text)

    def test_flat_archives_preserve_every_document(self):
        for slug in CASES:
            case = ROOT / 'testakten' / slug
            sources = {p.name: p for p in case.iterdir() if p.is_file() and p.name[:2].isdigit()}
            expected = set(expected_arcnames(case))
            self.assertEqual(len(expected), 8)
            for filename in ('build-testakten-release-zips.py', 'build-testakten-einzelpdf-zips.py'):
                with self.subTest(case=slug, builder=filename), tempfile.TemporaryDirectory() as directory:
                    path, _ = builder(filename).build_single(case, Path(directory))
                    with zipfile.ZipFile(path) as archive:
                        names = archive.namelist()
                        self.assertEqual(len(names), len(set(names)))
                        self.assertTrue(all('/' not in n and chr(92) not in n for n in names))
                        self.assertFalse(any(n.endswith('.md') for n in names))
                        self.assertEqual(archive.read('README.txt'), NOTICE_BYTES)
                        self.assertIsNone(archive.testzip())
                        if 'einzelpdf' in filename:
                            self.assertEqual(set(names), expected | {'README.txt'})
                            for name in expected:
                                data = archive.read(name)
                                self.assertEqual(pdf_content_errors(data), [], name)
                                self.assertGreater(len(PdfReader(io.BytesIO(data)).pages), 0)
                        else:
                            self.assertEqual(set(names), set(sources) | {'README.txt', f'{slug}_gesamt.pdf'})
                            for name, source in sources.items():
                                self.assertEqual(archive.read(name), source.read_bytes())

    def test_case_navigation_and_german_description(self):
        for path in (PLUGIN / 'README.md', ROOT / 'testakten/README.md'):
            text = path.read_text(encoding='utf-8')
            for slug in CASES:
                self.assertIn(slug, text, path)
        manifest = json.loads((PLUGIN / '.claude-plugin/plugin.json').read_text())
        display = builder('inject-direkt-loslegen-section.py').markdown_text(manifest['description'])
        self.assertIn('Geldwäscheprüfung', display)
        self.assertIn('Fachabläufe', display)


if __name__ == '__main__':
    unittest.main()
