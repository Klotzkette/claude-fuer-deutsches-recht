#!/usr/bin/env python3
"""Build only the ten source records of the Kassel recruiting case."""

import argparse
import csv
import json
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "scripts/fixtures/ki-verordnung-hochrisiko-pruefer/case.json"
CASE = ROOT / "testakten/ki-hochrisiko-bewerbungsauswahl-kassel"


def write_docx(record, reference, target):
    doc = Document()
    doc.core_properties.author = record["author"]
    doc.core_properties.title = record["title"]
    doc.core_properties.subject = reference
    doc.core_properties.created = datetime(2026, 9, 28)
    doc.core_properties.modified = datetime(2026, 9, 28)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in (("Normal", 11), ("Title", 18), ("Heading 1", 13)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "Times New Roman")
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.08
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    section.header.paragraphs[0].text = f'{record["author"]} | {reference}'
    for run in section.header.paragraphs[0].runs:
        run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.add_run(f'{reference} | Seite ')
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(f' von {len(record["pages"])}')
    for index, page in enumerate(record["pages"]):
        if index:
            doc.add_page_break()
        else:
            doc.add_paragraph(record["title"], "Title")
            doc.add_paragraph(f'{record["date_label"]}\nAn: {record["recipient"]}')
        doc.add_paragraph(page["heading"], "Heading 1")
        for text in page["paragraphs"]:
            paragraph = doc.add_paragraph(text)
            paragraph.paragraph_format.widow_control = True
    doc.save(target)


def write_pdf(record, reference, target):
    body = ParagraphStyle("Body", fontName="Times-Roman", fontSize=11,
                          leading=14, spaceAfter=10, alignment=TA_LEFT)
    heading = ParagraphStyle("Heading", parent=body, fontName="Times-Bold",
                             fontSize=13, leading=16, spaceBefore=9, spaceAfter=12)
    title = ParagraphStyle("Title", parent=heading, fontSize=18, leading=21)

    def page_frame(canvas, document):
        canvas.saveState()
        canvas.setFillColor(colors.black)
        canvas.setFont("Times-Roman", 9)
        canvas.drawString(62, 809, f'{record["author"]} | {reference}')
        canvas.drawString(62, 32, f'{reference} | Seite {document.page} von {len(record["pages"])}')
        canvas.restoreState()

    document = SimpleDocTemplate(str(target), pagesize=A4, leftMargin=62,
                                 rightMargin=62, topMargin=57, bottomMargin=54,
                                 title=record["title"], author=record["author"])
    flow = []
    for index, page in enumerate(record["pages"]):
        if index:
            flow.append(PageBreak())
        else:
            flow.extend([Paragraph(escape(record["title"]), title),
                         Paragraph(escape(record["date_label"]), body),
                         Spacer(1, 3)])
        flow.append(Paragraph(escape(page["heading"]), heading))
        flow.extend(Paragraph(escape(text), body) for text in page["paragraphs"])
    document.build(flow, onFirstPage=page_frame, onLaterPages=page_frame)


def build(output):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    output.mkdir(parents=True, exist_ok=True)
    for record in data["documents"]:
        target = output / record["file"]
        if target.suffix == ".eml":
            message = EmailMessage(policy=SMTP)
            message["From"] = record["from"]
            message["To"] = record["to"]
            message["Date"] = format_datetime(datetime.fromisoformat(record["date"]))
            message["Subject"] = record["subject"]
            message["Message-ID"] = f'<{target.stem}@akte.fuldaform.example>'
            message.set_content("\n\n".join(record["paragraphs"]), charset="utf-8")
            # Python releases differ in trailing whitespace on folded headers.
            headers, body = message.as_bytes().split(b"\r\n\r\n", 1)
            headers = b"\r\n".join(line.rstrip(b" \t") for line in headers.split(b"\r\n"))
            target.write_bytes(headers + b"\r\n\r\n" + body)
        elif target.suffix == ".txt":
            header = f'{record["title"]}\n{data["reference"]}\n{record["author"]}\n{record["date_label"]}'
            target.write_text(header + "\n\n" + "\n\n".join(record["paragraphs"]) + "\n", encoding="utf-8")
        elif target.suffix == ".csv":
            with target.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream, delimiter=";")
                writer.writerow(record["columns"])
                writer.writerows(record["rows"])
        elif target.suffix == ".docx":
            write_docx(record, data["reference"], target)
        elif target.suffix == ".pdf":
            write_pdf(record, data["reference"], target)
        else:
            raise ValueError(f"Unsupported source format: {target.suffix}")
        print(target)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=CASE)
    build(parser.parse_args().output_dir)
