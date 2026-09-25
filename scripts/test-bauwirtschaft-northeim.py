#!/usr/bin/env python3
"""Fallbezogene Regressionen ohne Live-Modelltest. Autor: Klotzkette."""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timedelta
from email import policy
from email.parser import BytesParser
from collections import Counter
import argparse
import hashlib
import importlib.util
import io
import json
import os
import re
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from office_process import run_office
from testakte_office_pdf import office_binary
from testakte_disclaimer import NOTICE_BYTES, pdf_content_errors

ROOT=Path(__file__).resolve().parents[1]
SLUG='bauwirtschaft-vergabeverfahren-feuerwehrhaus-northeim'
CASE=ROOT/'testakten'/SLUG
ASSETS=None

def module(name,file):
    spec=importlib.util.spec_from_file_location(name,ROOT/'scripts'/file)
    m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m

def text(file):
    if file.suffix=='.docx':return '\n'.join(p.text for p in Document(file).paragraphs)
    if file.suffix=='.pdf':return '\n'.join(p.extract_text() or '' for p in PdfReader(file).pages)
    if file.suffix=='.eml':return BytesParser(policy=policy.default).parsebytes(file.read_bytes()).get_body(preferencelist=('plain',)).get_content()
    return file.read_text(encoding='utf-8-sig')

def cents(n):return Decimal(str(n)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)

class CaseTests(unittest.TestCase):
    def test_01_scope_and_file_count(self):
        originals=[p for p in CASE.iterdir() if p.name[:2].isdigit()]
        self.assertEqual(len(originals),32)
        self.assertEqual(sorted(int(p.name[:2]) for p in originals),list(range(1,33)))
        self.assertEqual(Counter(p.suffix for p in originals),{'.docx':10,'.pdf':9,'.xlsx':5,'.eml':4,'.png':2,'.csv':1,'.txt':1})
        self.assertFalse(list(CASE.rglob('*.ndjson')))
        self.assertFalse(list(CASE.rglob('*.json')))
        self.assertFalse(list(CASE.rglob('*.pyc')))
        self.assertLessEqual(len(SLUG),64)
        for plugin in ('bauwirtschaft','fachanwalt-vergaberecht'):
            manifest=ROOT/plugin/'.claude-plugin/plugin.json'
            self.assertEqual(json.loads(manifest.read_text())['name'],plugin)
        for p in [ROOT/'scripts'/n for n in ('build-bauwirtschaft-northeim-akte.py','build-bauwirtschaft-northeim-workbooks.mjs','build-bauwirtschaft-northeim-pakete.py')]:
            self.assertNotIn('/Users/',p.read_text())

    def test_02_amounts_formulas_and_position_text(self):
        a=module('northeim_data','build-bauwirtschaft-northeim-akte.py')
        self.assertEqual(sum([820000,2120000,1140000,1690000,760000,720000,380000,770000]),8400000)
        self.assertGreater(8400000,5404000)
        for index,(prefix,net) in enumerate([('04',None),('08',313600),('10',356178),('12',371876)]):
            file=next(CASE.glob(prefix+'*.xlsx'))
            values=load_workbook(file,data_only=True);formulas=load_workbook(file,data_only=False)
            s=values.worksheets[0];lv=values['LV'];lf=formulas['LV']
            expected=Decimal(0)
            for i,p in enumerate(a.POSITIONS,6):
                self.assertEqual(lv[f'A{i}'].value,p[0]);self.assertEqual(lv[f'A{i}'].data_type,'s')
                self.assertEqual(lv[f'A{i}'].number_format,'@')
                self.assertEqual(lv[f'C{i}'].value,p[2]);self.assertEqual(lv[f'D{i}'].value,p[3])
                self.assertEqual(lv[f'B{i}'].value,p[1]+'. '+p[4])
                self.assertEqual(lf[f'F{i}'].data_type,'f')
                if net is not None:
                    expected+=cents(Decimal(str(p[2]))*Decimal(str(p[4+index])))
                    self.assertEqual(cents(lv[f'F{i}'].value),cents(p[2]*p[4+index]))
            if net is None:self.assertIn(s['B11'].value,(None,''))
            else:
                self.assertEqual(expected,Decimal(net))
                self.assertEqual(cents(s['B11'].value),Decimal(net))
                self.assertEqual(cents(s['B13'].value),cents(Decimal(net)*Decimal('.19')))
                self.assertEqual(cents(s['B14'].value),cents(Decimal(net)*Decimal('1.19')))
            self.assertEqual(formulas.properties.creator,'Klotzkette')
            for sh in formulas:
                for row in sh:
                    for cell in row:
                        self.assertNotEqual(cell.data_type,'e',(file.name,sh.title,cell.coordinate,cell.value))
            values.close();formulas.close()
        w=load_workbook(CASE/'23_preisspiegel.xlsx',data_only=True)
        self.assertEqual([w.active[f'{c}27'].value for c in 'BCD'],[313600,356178,371876])
        self.assertEqual([w.active[f'{c}29'].value for c in 'BCD'],[1,2,3]);w.close()

    def test_03_dates_and_transmissions(self):
        import csv
        with (CASE/'17_versandprotokoll.csv').open(encoding='utf-8-sig') as stream:rows=list(csv.DictReader(stream))
        events={(r['Kennung'],r['Status']):datetime.strptime(r['Zeit MESZ'],'%d.%m.%Y %H:%M:%S') for r in rows}
        sent=events['NO26-I025','elektronisch abgesendet'];earliest=sent.replace(hour=0,minute=0,second=0)+timedelta(days=11)
        self.assertEqual(earliest,datetime(2026,9,25));self.assertEqual(earliest.weekday(),4)
        self.assertEqual((earliest-sent.replace(hour=0,minute=0,second=0)).days,11)
        request=events['LT-NO26-NP01','Empfangseinrichtung gespeichert']
        ruege=events['NO26-R027','SMTP angenommen'];reply=events['NO26-S028','versandt und zugestellt']
        self.assertLess(request,earliest);self.assertLessEqual((ruege-sent).days,10);self.assertLessEqual((request-reply).days,15)
        self.assertLess(events['NO26-A019','gespeichert'],datetime(2026,9,9,12))
        for ident in ('NO26-E001','NO26-E002','NO26-E003'):self.assertLess(events[ident,'gespeichert'],datetime(2026,8,31,10))
        for file in CASE.glob('2[56]*.pdf'):
            t=text(file);self.assertIn('14.09.2026',t);self.assertIn('25.09.2026',t);self.assertIn('24.09.2026',t)
        t=text(CASE/'32_sachstand_25092026.docx')
        self.assertIn('keinen Zuschlag',t);self.assertIn('keinen Beschluss',t);self.assertIn('30.09.2026, 12:00',t)

    def test_04_mime_and_authors(self):
        for f in CASE.glob('*.eml'):
            raw=f.read_bytes();m=BytesParser(policy=policy.default).parsebytes(raw)
            self.assertFalse(m.defects);self.assertEqual(m['MIME-Version'],'1.0')
            self.assertEqual(m.get_content_type(),'multipart/alternative')
            for h in ('From','To','Date','Subject','Message-ID','Received'):self.assertIsNotNone(m[h])
            self.assertNotIn(b'\n',raw.replace(b'\r\n',b''))
            parts=[p for p in m.walk() if p.get_content_maintype()=='text']
            self.assertEqual(len(parts),2)
            for p in parts:self.assertEqual(p.get_content_charset(),'utf-8')
        for f in CASE.glob('*.docx'):
            d=Document(f);self.assertEqual(d.core_properties.author,'Klotzkette');self.assertEqual(d.styles['Normal'].font.name,'Times New Roman');self.assertEqual(d.styles['Normal'].font.size.pt,11)
        for f in CASE.glob('*.pdf'):self.assertEqual(PdfReader(f).metadata.author,'Klotzkette')

    def test_05_send_hashes_and_no_solution(self):
        receipt=text(CASE/'30_uebermittlungsbeleg.pdf')
        names=re.findall(r'\b\d{2}_[a-z0-9_]+\.(?:pdf|docx|xlsx|png|eml)',receipt)
        unique=set(names);self.assertEqual(len(unique),16)
        for n in unique:self.assertIn(hashlib.sha256((CASE/n).read_bytes()).hexdigest(),receipt)
        for f in CASE.iterdir():
            if f.suffix not in {'.pdf','.docx','.eml','.txt'}:continue
            t=text(f)
            for marker in ('Diese Testakte','This test case','Lösungsskizze','Musterlösung','Ausbildungsauftrag','§','TODO'):
                self.assertNotIn(marker,t,f.name)

    def test_06_repository_total(self):
        file=CASE/'gesamt-pdf'/f'{SLUG}_gesamt.pdf'
        self.assertGreaterEqual(len(PdfReader(file).pages),32)
        self.assertFalse(pdf_content_errors(file.read_bytes()))

    def test_08_archives(self):
        if ASSETS is None:self.skipTest('Exportprüfung nur mit --assets DIR')
        rz=module('northeim_rz','validate-testakten-release-zips.py')
        ez=module('northeim_ez','validate-testakten-einzelpdf-zips.py')
        source=ASSETS/f'testakte-{SLUG}.zip'
        single_name=f'testakte-{SLUG}-einzelpdfs.zip'
        candidates=[ASSETS/'testakten-einzelpdfs'/single_name,ASSETS/single_name]
        single=next((p for p in candidates if p.is_file()),candidates[0])
        self.assertTrue(source.is_file(),f'Originalarchiv fehlt: {source}')
        self.assertTrue(single.is_file(),f'Einzel-PDF-Archiv fehlt in {ASSETS} oder testakten-einzelpdfs')
        rz.assert_same('Originale',rz.expected_entries(CASE),rz.zip_entries(source,require_notice=True))
        ez.assert_same('Einzel-PDF',[ez.NOTICE_FILENAME,*ez.expected_arcnames(CASE)],ez.zip_entries(single,expected_suffix='.pdf'))
        with zipfile.ZipFile(single) as z:
            self.assertEqual(len(z.namelist()),33)
            for n in z.namelist():
                if n=='README.txt':self.assertTrue(z.read(n).startswith(NOTICE_BYTES));continue
                reader=PdfReader(io.BytesIO(z.read(n)))
                if n.startswith(('04_','08_','10_','12_')):
                    t='\n'.join(p.extract_text() or '' for p in reader.pages)
                    for p in ('01.010','01.100','01.200'):self.assertIn(p,t)
                if n.startswith('23_'):
                    compact=re.sub(r'\s+','','\n'.join(p.extract_text() or '' for p in reader.pages))
                    for amount in ('313.600,00','356.178,00','371.876,00'):self.assertIn(amount,compact)
        with zipfile.ZipFile(source) as z:
            self.assertEqual(len(z.namelist()),34)
            self.assertTrue(z.read('README.txt').startswith(NOTICE_BYTES))
            self.assertFalse(any(n.endswith(('.json','.ndjson','.yaml','.md')) for n in z.namelist()))
            for original in CASE.iterdir():
                if original.is_file() and original.name[:2].isdigit():
                    self.assertEqual(z.read(original.name),original.read_bytes(),original.name)
            total=PdfReader(io.BytesIO(z.read(f'{SLUG}_gesamt.pdf')))
            self.assertGreaterEqual(len(total.pages),32)
            content='\n'.join(p.extract_text() or '' for p in total.pages)
            for marker in ('NO-FH26-L430','Nachprüfungsantrag','25.09.2026','313.600,00','356.178,00','371.876,00'):
                self.assertIn(marker,content)

    def test_07_readme_and_strict_validator(self):
        r=module('northeim_readme','validate-testakten-readme-downloads.py');errors=[]
        r.validate_local_readme(SLUG,CASE,errors);self.assertEqual(errors,[])
        strict=module('northeim_strict','validate-testakten-dokumentqualitaet.py')
        with tempfile.TemporaryDirectory(prefix='northeim-strict-') as temp:
            view=Path(temp);(view/SLUG).symlink_to(CASE,target_is_directory=True)
            strict.TESTAKTEN=view;strict.REPO=view
            self.assertEqual(strict.main(),0)

def native_regression():
    binary=office_binary()
    if not binary:raise RuntimeError('SOFFICE für --native erforderlich')
    source=CASE/'08_angebot_leinetal.xlsx';before=hashlib.sha256(source.read_bytes()).hexdigest()
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda n:f'{{{ns}}}{n}'
    with tempfile.TemporaryDirectory(prefix='northeim-native-input-') as tmp:
        temp=Path(tmp);output=temp/'out';output.mkdir()
        cases=[('preis',78123.45,313723.45),('nullpreis',0,235600),('leer',None,None)]
        for name,value,expected in cases:
            with zipfile.ZipFile(source) as src,zipfile.ZipFile(temp/(name+'.xlsx'),'w',zipfile.ZIP_DEFLATED) as dest:
                for info in src.infolist():
                    data=src.read(info.filename)
                    if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',info.filename):
                        root=ET.fromstring(data)
                        if info.filename=='xl/worksheets/sheet2.xml':
                            cell=next(c for c in root.iter(q('c')) if c.get('r')=='E7')
                            for e in list(cell):cell.remove(e)
                            cell.attrib.pop('t',None)
                            if value is not None:ET.SubElement(cell,q('v')).text=str(value)
                        for cell in root.iter(q('c')):
                            if cell.find(q('f')) is not None:
                                for cached in list(cell):
                                    if cached.tag in (q('v'),q('is')):cell.remove(cached)
                                cell.attrib.pop('t',None)
                        data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
                    dest.writestr(info,data)
        cmd=[binary,f'-env:UserInstallation={(temp/"profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(output),*[str(temp/(name+'.xlsx')) for name,_,_ in cases]]
        code,details=run_office(cmd,timeout=180,env=os.environ.copy());assert code==0,details
        for name,value,expected in cases:
            w=load_workbook(output/(name+'.xlsx'),data_only=True);actual=w['Angebot']['B11'].value
            if expected is None:assert actual in (None,''),(name,actual)
            else:assert cents(actual)==cents(expected),(name,actual,expected)
            w.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before
    print('Native LibreOffice-Eingabevarianten OK: Preisänderung, Nullpreis, Leerpreis; Original unverändert.')

if __name__=='__main__':
    parser=argparse.ArgumentParser(description='Standardlauf: ausschließlich Repository-Prüfungen.')
    parser.add_argument('--native',action='store_true',help='Zusätzliche LibreOffice-Neuberechnung in temporären Kopien')
    parser.add_argument('--assets',type=Path,help='Zusätzliche Archivprüfung; Einzel-PDF-ZIP hier oder unter testakten-einzelpdfs')
    args=parser.parse_args()
    ASSETS=args.assets.resolve() if args.assets else None
    result=unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(CaseTests))
    if not result.wasSuccessful():raise SystemExit(1)
    if args.native:native_regression()
