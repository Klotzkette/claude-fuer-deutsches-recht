#!/usr/bin/env python3
"""Reproduzierbare native Arbeitsakte Auenhöhe; erstellt ausschließlich 01–36."""
from pathlib import Path
from datetime import datetime, timezone
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from xml.sax.saxutils import escape
import json, re, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from akten_build_runtime import serif_font_path
ROOT=Path(__file__).resolve().parents[1]
CASE=ROOT/'testakten/krankenhaus-it-ki-auenhoehe-thueringen'
CASE.mkdir(parents=True,exist_ok=True)
STAMP=datetime(2026,10,6,14,0,tzinfo=timezone.utc)
H='Klinikverbund Auenhöhe GmbH | Lindenried bei Saalfeld'
DOCS=[]
def add(n,stem,kind,title,author,date,text,to='Nora Bergmann | Gesamtleitung IT'):
    DOCS.append(dict(n=n,name=f'{n:02d}_{stem}.{kind}',kind=kind,title=title,author=author,date=date,text=text.strip(),to=to))
add(1,'Projektauftrag','docx','Auftrag zur Vorbereitung der drei Digitalisierungsvorhaben','Dr. Selma Hartung | Geschäftsführung','21.09.2026', '''1 Ziel und Verantwortung

Frau Bergmann, bereiten Sie bitte eine entscheidungsfähige Vorlage für die Digitalisierung unserer Dokumentation, für die radiologische Priorisierung und für die retrospektive Versorgungsforschung vor. Die Projekte heißen P01, P02 und P03. Ich möchte am 9. Oktober erkennen können, welche Schritte wir verantworten können und welche Unterlagen noch fehlen. Ein gemeinsamer Projektordner soll die Abstimmung erleichtern; er ist keine gemeinsame Genehmigung aller Verarbeitungen.

2 Vorgesehener Pilot

Für P01 wünschen wir einen achtwöchigen Pilot ab dem 19. Oktober 2026. Der Termin ist ein Planungsziel. Eine Anbindung an das produktive Krankenhausinformationssystem, die Übermittlung von Patientendaten und der Einsatz im Behandlungsablauf sind damit nicht freigegeben. Bitte beginnen Sie die technischen Prüfungen mit eigens geschriebenen Beispielsätzen und einer getrennten Testumgebung. Eine bloße Demonstration beim Anbieter ersetzt keine Freigabe durch unser Haus.

3 Beteiligte und Entscheidungen

Sie führen das Vorhaben aus Sicht der IT. Herr Fink begleitet es als Datenschutzbeauftragter; Herr Seidel verantwortet die technische Sicherheitsprüfung. Frau Dr. Feld klärt die ärztlichen Anforderungen. Die Entscheidung über Budget und Einsatzumfang bleibt bei der Geschäftsführung. Die ärztliche Leitung entscheidet nicht allein über Datenübermittlungen. Der Betriebsrat erhält die Funktionen, die Beschäftigte betreffen, vor einer Einführung zur Beratung.

4 Wirtschaftliche Vorbereitung

Die Beschaffung soll zunächst auf die acht Wochen begrenzt werden. Eine automatische Verlängerung möchte ich nicht ohne erneute Vorlage. Zeitersparnisse sind erst dann belastbar, wenn die Zeit für Kontrolle und Korrektur mitgemessen wird. Bitte weisen Sie laufende Kosten, Einrichtung und internen Aufwand getrennt aus. Ein Preisangebot ist noch keine Bestellung. Für P02 und P03 erwarte ich zunächst eine gesonderte Aufstellung ihrer Voraussetzungen.

Dr. Selma Hartung''')
add(2,'Betriebsprofil','docx','Struktur des Klinikverbunds und Zuständigkeiten','Nora Bergmann | Gesamtleitung IT','22.09.2026','''1 Haus und Versorgungsauftrag

Der Klinikverbund Auenhöhe GmbH betreibt in Lindenried bei Saalfeld ein Krankenhaus mit 240 Betten. Im Jahr 2025 wurden rund 14.500 vollstationäre Fälle behandelt. Das Haus verfügt unter anderem über Innere Medizin, Allgemeinchirurgie, Intensivmedizin und eine leistungsfähige Radiologie. Die alleinige kommunale Gesellschafterin ist der Landkreis Saalebogen. Der Standort ist ländlich geprägt; Fachabteilungen arbeiten mit Einweisern und ambulanten Nachbehandlern zusammen.

2 Medizinisches Versorgungszentrum

Die Auenhöhe MVZ GmbH ist eine vollständig gehaltene Tochtergesellschaft mit eigener Geschäftsführung und eigener Patientenverwaltung. Gemeinsame Personalabteilung und zentrale IT bedeuten derzeit keine gemeinsame Patientenakte. Das MVZ nutzt ein getrenntes System. Eine Zuordnung von Krankenhausfällen zu späteren ambulanten Behandlungen ist technisch nur mit einem zusätzlichen Abgleich möglich. Die Beteiligten haben hierfür noch keinen Datenübermittlungsvertrag unterzeichnet.

3 Ansprechpartner

Die Geschäftsführerin des Krankenhauses ist Dr. Selma Hartung. Nora Bergmann leitet die IT und koordiniert die drei Vorhaben. Benedikt Fink ist Datenschutzbeauftragter, Karim Seidel Informationssicherheitsbeauftragter. Chefärztin Dr. Amira Feld vertritt die klinischen Nutzer. Dr. Oskar Wenck koordiniert die Versorgungsforschung. Paul Lindner ist Vorsitzender des Betriebsrats. Im MVZ ist Geschäftsführerin Ruth Demir die Ansprechpartnerin für Datenbereitstellungen.

4 Bestehende Systeme

Das Krankenhausinformationssystem liegt im eigenen Rechenraum mit täglicher Sicherung an einen zweiten Standort. Das Bildarchiv der Radiologie ist ein gesondertes System. Beschäftigte erhalten persönliche Zugänge; einzelne Funktionspostfächer werden von wechselnden Mitarbeitenden gelesen. Der vorhandene Fernwartungszugang wird zentral freigeschaltet. Seine Rechte reichen derzeit weiter als für P01 vorgesehen. Die IT zählt neun Beschäftigte einschließlich Leitung, Teilzeit und Servicekoordination.

Nora Bergmann''')
add(3,'Geschaeftsfuehrung_Starttermin','eml','Bitte am Freitag eine belastbare Entscheidungsvorlage','Dr. Selma Hartung <selma.hartung@auenhoehe.example>','2026-09-23T08:17:00','''Guten Morgen Nora,

ich habe gestern im Aufsichtsrat von den drei Vorhaben berichtet. Dort kam verständlicherweise sofort die Frage nach dem Start. Ich habe den 19. Oktober als Wunsch für P01 genannt, nicht als zugesagten Produktivtermin. Bitte lass diese Unterscheidung auch in der Vorlage stehen.

Wenn die Vertragsunterlagen noch nicht vollständig sind, brauche ich eine konkrete Beschreibung dessen, was bis dahin geprüft werden kann. Für eine Vorführung dürfen wir keine Patientenakte aus dem Stationssystem nehmen. Die Ärztinnen können eigene Beispieldiktate erstellen; dafür brauchen wir keine echten Namen oder Aufnahmebefunde.

Bitte rechne die interne Kontrollzeit ehrlich mit. Wir gewinnen nichts, wenn ein Brief in einer Minute entsteht und danach zwanzig Minuten geprüft werden muss. Der achtwöchige Pilot soll nach einer Entscheidung starten, nicht einfach durch Ablauf unserer internen Frist. Für Radiologie und Forschung bitte jeweils einen eigenen Abschnitt.

Viele Grüße
Selma''')
add(4,'Angebot_Sprechfeder','pdf','Angebot für den befristeten Einsatz von Sprechfeder Clinical 2.4','Sprechfeder Health Europe GmbH | Maja Roth','24.09.2026','''Sehr geehrte Frau Bergmann,

wir bieten Ihnen für P01 dreißig persönliche Nutzerzugänge zu Sprechfeder Clinical 2.4 für zwei Monate an. Der Preis beträgt 89,00 EUR netto je Nutzer und Monat. Daraus ergeben sich 5.340,00 EUR netto für die Zugänge. Einrichtung, Schulung und die Konfiguration einer Testanbindung berechnen wir mit einmalig 4.800,00 EUR netto. Das Angebot beträgt insgesamt 10.140,00 EUR netto; die gesetzliche Umsatzsteuer wird zusätzlich berechnet. Ein Erwerb klinischer Entscheidungsfunktionen ist nicht Gegenstand dieses Angebots.

Die Anwendung transkribiert Diktate und schlägt auf Basis der vom Nutzer eingegebenen Inhalte einen Briefentwurf vor. Jede Ausgabe bleibt bis zur Übernahme durch einen berechtigten Nutzer ein Entwurf. Eine automatisierte Übernahme in die freigegebene Patientenakte ist nicht Teil der angebotenen Konfiguration. Schnittstellenänderungen am Krankenhausinformationssystem sind nicht enthalten.

Die Verarbeitungsregion ist Frankfurt am Main. Support durch unser US-Unternehmen kann für schwierige Fehleranalysen angefordert werden. Nach unserem derzeitigen Standardvertrag dürfen in diesem Zusammenhang notwendige Diagnoseinformationen verarbeitet werden. Welche Daten dies im Einzelfall umfasst, stimmen wir im Einrichtungsgespräch ab. Die Geschäftsbedingungen enthalten außerdem eine Regelung zur Verbesserung unserer Modelle; Ihr Änderungswunsch hierzu ist noch in Bearbeitung.

Der Leistungszeitraum beginnt mit der schriftlichen Abrufbestätigung. Für die zwei Monate bieten wir die genannten Konditionen bis zum 12. Oktober 2026 an. Über eine Fortführung ist gesondert zu verhandeln. Die Unterzeichnung dieses Angebots ersetzt nicht die Abstimmung des Datenschutzanhangs oder Ihrer internen Nutzungsregeln.

Mit freundlichen Grüßen
Maja Roth | Kundenbetreuung''')
add(5,'AVV_Lieferantenentwurf','docx','Vereinbarung zur Verarbeitung im Auftrag Lieferantenentwurf','Sprechfeder Health Europe GmbH | Vertragsabteilung','24.09.2026','''1 Parteien und Gegenstand

Die Sprechfeder Health Europe GmbH verarbeitet für die Klinikverbund Auenhöhe GmbH die vom Auftraggeber bereitgestellten Diktate, Transkripte und Briefentwürfe innerhalb der Anwendung Sprechfeder Clinical 2.4. Die Verarbeitung dient der Transkription und der Erstellung von Entwürfen. Der Auftraggeber bestimmt, welche Beschäftigten Inhalte einstellen und welche Ausgaben übernommen werden. Der Entwurf ist noch nicht unterzeichnet.

2 Daten und Verarbeitung

Die bereitgestellten Inhalte können Identitätsdaten, Kontaktdaten, Gesundheitsangaben und Angaben über behandelnde Beschäftigte enthalten. Sprachaufnahmen werden nach unserem Standardangebot sieben Tage, gespeicherte Entwürfe dreißig Tage aufbewahrt. Der Auftraggeber kann die Entwürfe vorher löschen. Betriebsprotokolle werden nach dem vorgesehenen Standard neunzig Tage aufbewahrt. Für Sicherungskopien gelten die in den technischen Maßnahmen beschriebenen Zyklen.

3 Weisungen und Unterstützung

Der Auftragnehmer verarbeitet die Inhalte auf dokumentierte Weisung des Auftraggebers und unterstützt ihn bei Auskunfts-, Berichtigungs- und Löschbegehren. Er meldet bekannt gewordene Verletzungen des Schutzes personenbezogener Daten unverzüglich an die benannte Kontaktstelle. Die inhaltliche Beantwortung von Patientenanfragen bleibt beim Auftraggeber. Der Auftragnehmer stellt Nachweise über die vereinbarten Maßnahmen zur Verfügung; die Einzelheiten einer Prüfung werden vorher abgestimmt.

4 Produktverbesserung nach Lieferantenstandard

Der Auftraggeber gestattet nach diesem Entwurf die Nutzung von Eingaben und korrigierten Ausgaben zur Verbesserung der Modelle des Auftragnehmers. Ein Widerspruch kann über die Administrationskonsole erklärt werden und wirkt nach dem Standardangebot für neue Verarbeitungsvorgänge. Diese vom Lieferanten vorgeschlagene Regelung ist Bestandteil des übersandten Entwurfs. Die Klinik hat ihr bislang nicht zugestimmt; der Auftragnehmer bearbeitet den Änderungswunsch aus dem Einkauf.

5 Unterauftragnehmer und Vertragsende

Der Auftragnehmer setzt die Unternehmen aus der gesondert übersandten Liste ein. Änderungen teilt er vor Umsetzung mit. Nach Vertragsende stellt er die gespeicherten Entwürfe zum Export bereit und löscht sie nach der vereinbarten Rückgabefrist. Der Umgang mit Modellen, Sicherungen und Diagnosepaketen wird im Standardtext nicht gesondert beschrieben. Die Parteien haben zu diesen Punkten noch keinen abschließenden Text vereinbart.

Vertragsabteilung Sprechfeder Health Europe GmbH''')
add(6,'TOM_und_Pruefbericht','pdf','Technische Maßnahmen und vorliegender Prüfbericht','Sprechfeder Health Europe GmbH | Felix Sommer','25.09.2026','''Sehr geehrter Herr Seidel,

die Kundendatenbank unseres europäischen Angebots liegt in Frankfurt am Main. Die Übertragung erfolgt verschlüsselt; die Speicherverschlüsselung wird durch die Plattform verwaltet. Die Schlüsselverwaltung liegt beim Betreiber der Plattform. Für Administrationszugriffe verwenden wir persönliche Konten und eine zusätzliche Authentisierung. Supportzugriffe werden in einem Ticket dokumentiert. Der Standardrollenplan unterscheidet Klinikadministration, Nutzer und Support.

Wir verweisen auf einen C5-Prüfbericht Typ 1 zum Stichtag 15. April 2026 für unsere Plattform Health Workspace. Der Bericht beschreibt die Ausgestaltung der dort aufgeführten Kontrollen zu diesem Zeitpunkt. Der hier übersandte Deckblattauszug benennt Health Workspace, jedoch nicht ausdrücklich Sprechfeder Clinical 2.4. Die Anlage mit ergänzenden Kontrollkriterien beim Kunden ist in dieser Lieferung nicht enthalten. Wir prüfen, ob der Bericht auch die Diagnosedatenverarbeitung durch den US-Support abdeckt.

Nach Auskunft unseres Produktteams wurde Health Workspace erstmals am 4. November 2025 in Verkehr gebracht. Die Belege zur Markteinführung und die Abgrenzung gegenüber einem früheren Entwicklungsdienst haben wir angefordert. Mit diesem Schreiben bestätigen wir weder die vollständige Systemabdeckung noch einen bestimmten gesetzlichen Erfüllungsstatus Ihres Einsatzes. Die gesonderte Prüfung für Ihr Krankenhaus bleibt damit offen.

Die tägliche Sicherung wird dreißig Tage vorgehalten. Für das Löschen einzelner Inhalte aus Sicherungen ist noch kein kundenseitiger Nachweisabruf vorgesehen. Inaktive Nutzerkonten können deaktiviert werden; die Aufbewahrung ihrer Protokolle folgt einem getrennten Ablauf. Bitte teilen Sie uns mit, welche Aufbewahrungs- und Exportanforderungen Ihre Klinik konkret benötigt. Eine Bestätigung über die Abschaltung des Modelltrainings liegt aus unserem Produktteam noch nicht vor.

Mit freundlichen Grüßen
Felix Sommer | Informationssicherheit''')
add(7,'Unterauftragsverarbeiter','docx','Liste der vorgesehenen Dienstleister und Zugriffe','Sprechfeder Health Europe GmbH | Maja Roth','25.09.2026','''1 Europäische Plattform

Die Sprechfeder Health Europe GmbH ist Vertragspartnerin des Krankenhauses. Für den Betrieb der Plattform wird MainCloud Services GmbH mit Rechenzentrumsregion Frankfurt eingesetzt. Dort liegen nach dem derzeitigen Betriebsmodell Datenbank, Objektspeicher und Sicherungen. Die technische Plattform stellt Speicher und Rechenleistung bereit. Eine eigene Nutzung von Krankenhausinhalten durch MainCloud ist in unserem Lieferantenvertrag nicht vorgesehen.

2 Support in den Vereinigten Staaten

Sprechfeder Support Inc. soll bei eskalierten Störungen unterstützen. Ihr Zugriff ist nach Vertriebsangabe auf die erforderlichen Diagnoseinformationen beschränkt. Die tatsächlichen Rechte des Supportkontos wurden für Ihren Mandanten noch nicht eingerichtet und geprüft. Bei Fehlern der Transkription können nach der Supportbeschreibung auch Eingaben oder Ausgaben in ein Diagnosepaket gelangen. Ein allgemeiner Ausschluss von Inhaltsdaten wird mit dieser Liste nicht erklärt.

3 Nachweise zum Übermittlungsweg

Wir haben die Nachweisführung für den Drittlandzugriff bei unserer Rechtsabteilung angefordert. Einen aktuellen Eintrag in einer Zertifizierungsliste werden wir gegebenenfalls nachreichen. Diese Ankündigung enthält noch keine Bestätigung, dass gerade die benannte juristische Person und der vorgesehene Datenumfang erfasst sind. Alternative Vertragsunterlagen sind mit diesem Schreiben ebenfalls noch nicht übermittelt.

4 Weitere Leistungen

Die europäischen Mitarbeiter nutzen ein Ticketsystem der Helpdesk Kontor GmbH. Das Ticket enthält Ansprechpartner, technische Fehlermeldung und gegebenenfalls eine beigefügte Diagnosedatei. Im Standardformular fehlt eine technische Sperre gegen das Anhängen von Patientendokumenten. Die für Ihre Installation verantwortlichen Beschäftigten werden im Einrichtungstermin benannt. Veränderungen an dieser Liste sollen Ihnen vor dem produktiven Einsatz mitgeteilt werden.

Maja Roth''')
add(8,'Lieferant_Angebot_Anhang','eml','Angebot P01 und noch offene Anlage zur Produktverbesserung','Maja Roth <maja.roth@sprechfeder.example>','2026-09-25T11:28:00','''Sehr geehrte Frau Bergmann,

anbei finden Sie unser Angebot vom 24. September. Die Preise gelten für die dreißig Nutzer und genau zwei Monate. Eine Einrichtungsleistung ist einmalig enthalten; Anpassungen Ihres Klinikinformationssystems wären gesondert zu besprechen.

Den Vertrag habe ich nicht als bereits abgestimmt markiert. Ihr Wunsch, jegliche Verwendung der Klinikdaten für unsere Modellverbesserung auszuschließen, liegt bei der Vertragsabteilung. Die Schaltfläche im Administrationsmenü reicht Ihnen nach unserem Telefonat nicht als Nachweis. Ich habe deshalb eine verbindliche Beschreibung der Einstellung und ihrer Wirkung auf Diagnosepakete angefordert.

Für einen aktuellen Nachweis zu Sprechfeder Support Inc. kann ich Ihnen heute noch keinen Link schicken. Bitte behandeln Sie meine Aussage aus der Vorführung, wir hätten dafür alle üblichen Unterlagen, nicht als Übergabe dieser Unterlagen. Herr Sommer meldet sich gesondert zur Plattformprüfung.

Freundliche Grüße
Maja Roth''')
add(9,'Datenflussaufnahme','docx','Aufnahme der Datenflüsse für die drei Vorhaben','Nora Bergmann und Karim Seidel','28.09.2026','''1 P01 Diktat und Briefentwurf

Ein berechtigter Beschäftigter soll eine Sprachaufnahme oder einen Text an Sprechfeder Clinical 2.4 übergeben. Der Browser überträgt die Eingabe an die Plattform in Frankfurt. Dort entstehen Transkript und Briefentwurf. Der Nutzer prüft die Ausgabe und übernimmt sie bei Bedarf manuell in das Krankenhausinformationssystem. Die vorgesehene Testkonfiguration schreibt nicht automatisch in die Patientenakte. Die Anbindung eines produktiven Patientenkontexts ist noch nicht eingerichtet.

Ein Fehler kann ein Ticket mit einem Diagnosepaket auslösen. Ob dieses Paket automatisch Ausschnitte aus Diktat oder Entwurf enthält, ist noch offen. Der europäische Support kann ein Ticket an die US-Gesellschaft eskalieren. Im Angebot steht Datenregion Frankfurt; die technische Supportbeschreibung enthält einen zusätzlichen Zugriffsweg. Wir haben diesen Weg noch nicht am eigenen Mandanten getestet.

2 P02 Bilddaten und Priorisierung

Für Radialert Triage 5.1 ist eine lokale Verarbeitung von Kopien ausgewählter Bildserien vorgesehen. Die Ausgabe soll eine zusätzliche Priorisierungsmarkierung in einer Arbeitsliste erzeugen. Das Bildarchiv bleibt das führende System. Die Zuordnung von Untersuchung und Ausgabe erfolgt über interne Kennungen. Ein Herstellerzugang zur Wartung soll zeitlich begrenzt geöffnet werden. Die Übernahme neuer Modellversionen ist derzeit nicht automatisiert vorgesehen.

3 P03 Forschung und Nachverfolgung

Das Krankenhaus soll stationäre Fälle aus den Jahren 2022 bis 2025 pseudonymisieren. Die Zuordnungstabelle bleibt beim Krankenhaus. Für die geplante Nachverfolgung wären zusätzlich ambulante Daten aus der MVZ GmbH erforderlich. Ihre Herkunft und der Abgleich mit dem Krankenhausbestand müssen gesondert festgelegt werden. Die Hochschule Saalebogen soll einen Datensatz ohne Namen erhalten. Ein vom Projektplan abweichendes Training eines allgemein einsetzbaren Modells hat das Forschungsteam zusätzlich angesprochen, aber noch nicht beschrieben.

4 Stand der Aufnahme

Diese Aufnahme beschreibt geplante Verbindungen. Sie bestätigt weder Rechtsgrundlagen noch eine Freigabe. Die tatsächlich übertragenen Felder, Supportrechte und Löschwege sind bei der technischen Erprobung zu protokollieren. Für jedes Vorhaben bleibt ein eigener Umfang erforderlich.

Nora Bergmann''')
add(10,'Supportprobe','pdf','Protokoll der Supportprobe mit einem Beispieldiktat','Karim Seidel | Informationssicherheit','29.09.2026','''Am 29. September wurde zwischen 10:10 und 10:42 Uhr eine Supportprobe in der getrennten Testumgebung durchgeführt. Nora Bergmann gab einen eigens geschriebenen Beispielsatz ein. Der Satz enthielt keine Informationen aus einer Patientenakte. Das zugehörige Ticket trug die Kennung SF-DEMO-018. Ziel war die Prüfung, welche Informationen der Support ohne zusätzliche Freigabe sehen kann.

Nach Erzeugung einer Fehlermeldung bot die Anwendung eine Schaltfläche zum Versenden eines Diagnosepakets an. Die Vorschau zeigte Browserkennung, Nutzerkennung, Zeitstempel und zwei Zeilen des eingegebenen Texts. Herr Sommer erklärte im Gespräch, dies erleichtere die Fehleranalyse. Eine Bestätigung, dass diese Ausschnitte für den Fehler technisch notwendig seien, wurde nicht abgegeben. Die Klinik hat das Paket nicht an den Support übermittelt.

Der europäische Support konnte über eine Bildschirmfreigabe die Vorschau sehen. Ein Zugriff der US-Gesellschaft fand während der Probe nicht statt. Ob dieselbe Rolle über den normalen Eskalationsweg einen vollständigen Entwurf abrufen kann, ließ sich an diesem Termin nicht feststellen. Der Anbieter möchte die Rechteübersicht nachreichen. Die Region des Datenbankspeichers war in der Verwaltungskonsole mit Frankfurt bezeichnet; über entfernte Einsichtnahmen sagte dieses Feld nichts aus.

Nach der Probe wurde das Beispieldiktat in der Anwendung gelöscht. Der Papierkorb zeigte anschließend keinen Eintrag. Ein Löschprotokoll für Sicherungen oder Diagnoseprotokolle wurde nicht erzeugt. Es gab keine Feststellung einer Verletzung von Patientendaten, weil ausschließlich neu geschriebene Testinhalte verwendet wurden. Aus dem Versuch folgt keine Aussage, dass vergleichbare Vorgänge im Echtbetrieb ohne weitere Prüfung eingesetzt werden können.

Aufgenommen von Karim Seidel. Nora Bergmann hat Ablauf und Zeitangaben am selben Tag bestätigt.''')
add(11,'TIA_Fragenstand','docx','Fragen zum möglichen Zugriff aus den Vereinigten Staaten','Benedikt Fink | Datenschutzbeauftragter','30.09.2026','''Sehr geehrte Frau Roth,

für die Bewertung des vorgesehenen Supportwegs benötigen wir konkrete Unterlagen. Die Angabe einer europäischen Speicherregion beschreibt nicht, wer von wo auf die Inhalte zugreifen kann. Bitte beantworten Sie die folgenden Fragen bezogen auf die tatsächlich einzusetzenden juristischen Personen und die Version Sprechfeder Clinical 2.4. Eine allgemeine Konzernbroschüre genügt uns zur Zuordnung nicht.

1 Beteiligte und Daten

Welche Gesellschaft erhält im normalen Betrieb und bei einer Eskalation Zugriff? Wer beschäftigt die zugriffsberechtigten Personen und in welchen Staaten arbeiten sie? Bitte unterscheiden Sie Zugangsdaten, technische Protokolle, Sprachaufnahmen und Briefentwürfe. Welche dieser Daten lassen sich vor einer Supportweitergabe entfernen, ohne dass die Diagnose unmöglich wird? Bitte fügen Sie einen beispielhaften, inhaltsleeren Datensatz und den Rollenplan bei.

2 Übermittlungsinstrument und praktische Prüfung

Auf welches Instrument stützen Sie einen erforderlichen Drittlandzugriff? Wenn Sie eine Zertifizierung anführen, benötigen wir einen aktuellen Nachweis mit Unternehmen, erfasstem Datenbereich und überprüfbarem Status. Falls Vertragsklauseln eingesetzt werden, bitten wir um die vollständige vereinbarte Fassung einschließlich Anlagen. Bitte erläutern Sie die Bewertung der tatsächlichen Zugriffsumstände und etwaiger ergänzender Schutzmaßnahmen. Ein nicht bestätigter Verweis auf einen späteren Listeneintrag ist noch kein vorliegender Nachweis.

3 Technische Begrenzung und weitere Verwendung

Kann die Klinik jeden Zugriff einzeln freischalten, zeitlich begrenzen und nachträglich nachvollziehen? Wer kann die verschlüsselten Inhalte entschlüsseln? Welche Sicherungs- und Protokollkopien entstehen und wann werden diese gelöscht? Erfasst ein Ausschluss der Modellverbesserung auch Tickets, Diagnosepakete und durch Nutzer korrigierte Ausgaben? Bitte nennen Sie die vertragliche Zusage und die technisch wirksame Konfiguration getrennt.

4 Weiteres Vorgehen

Wir warten mit unserer abschließenden Bewertung auf Ihre Antwort. Bitte senden Sie die Unterlagen bis zum 7. Oktober an die Projektleitung und an mich. Ein laufender Fragebogen ersetzt keine abgeschlossene Transferbewertung.

Mit freundlichen Grüßen
Benedikt Fink''')
add(12,'DSB_Nachweise','eml','P01 Nachweise sind noch nicht vollständig','Benedikt Fink <benedikt.fink@auenhoehe.example>','2026-09-30T16:04:00','''Hallo Nora,

den Fragenstand an den Anbieter habe ich abgelegt. Wir sollten am Freitag sauber zwischen vorliegenden Dokumenten, Ankündigungen und einer abgeschlossenen Bewertung unterscheiden. Zur US-Gesellschaft fehlt uns ein überprüfbarer aktueller Nachweis; zur Modellverbesserung gibt es bislang nur den offenen Änderungswunsch.

Der C5-Auszug ist ein weiterer eigener Punkt. Herr Sommer nennt einen Prüfbericht Typ 1 und einen Zeitpunkt des erstmaligen Inverkehrbringens. Beides müssen wir mit dem tatsächlich eingesetzten System und dem relevanten Umfang verbinden. Aus dem Berichtstitel allein kann ich weder die Einbeziehung des Supportwegs noch die Erfüllung der vom Kunden zu leistenden Kontrollen ableiten.

Ich schlage vor, dass wir unsere offenen Fragen in der Maßnahmenmappe jeweils einer Person zuordnen. Für die Geschäftsführung kann ich den Stand erklären. Eine unterschriebene Betriebsfreigabe liegt von mir nicht vor; die Entscheidung der Geschäftsführung benötigt außerdem die technischen und klinischen Beiträge.

Viele Grüße
Benedikt''')
add(13,'Patienteninformation_Entwurf','docx','Patienteninformation zur unterstützten Erstellung von Briefentwürfen','Klinikverbund Auenhöhe GmbH | Kommunikation','30.09.2026','''Entwurf zur internen Abstimmung vom 30. September 2026. Dieser Text ist noch nicht veröffentlicht.

1 Geplante Unterstützung bei Arztbriefen

Unsere Beschäftigten sollen bei der Erstellung von Entlassbriefen eine Anwendung zur Transkription von Diktaten und zur Formulierung eines ersten Textentwurfs nutzen können. Die behandelnden Beschäftigten prüfen und bearbeiten den Entwurf. Der verantwortliche ärztliche Dienst entscheidet über den endgültigen Brief. Die Anwendung ist in diesem Vorhaben nicht dafür vorgesehen, selbst eine Diagnose zu stellen oder eine Behandlung festzulegen.

2 Verantwortliche Stelle und Fragen

Für die Verarbeitung im Krankenhaus ist die Klinikverbund Auenhöhe GmbH verantwortlich. Bei Fragen zur vorgesehenen Verarbeitung können Sie sich an die Patientenverwaltung oder an unseren Datenschutzbeauftragten Benedikt Fink unter datenschutz@auenhoehe.example wenden. Behandlungen im rechtlich eigenständigen Auenhöhe MVZ sind von dieser Information nicht automatisch erfasst.

3 Vorgesehene Daten und Dienstleister

Für den Entwurf können Diktate, Angaben zur Behandlung und die für den Brief erforderlichen Identitätsangaben verarbeitet werden. Nach derzeitiger Planung arbeitet das Krankenhaus mit der Sprechfeder Health Europe GmbH zusammen. Deren Plattform wird in Frankfurt betrieben. Ein möglicher Supportzugriff aus einem weiteren Staat wird noch geprüft. Die Klinik hat die endgültige Vereinbarung hierzu noch nicht geschlossen.

4 Noch abzustimmende Angaben

Die Kommunikationsabteilung hat die abschließenden Angaben zu Rechtsgrundlagen, Empfängern, Speicherfristen und Betroffenenrechten bei Herrn Fink angefordert. Die Klinik beabsichtigt nicht, mit dieser Entwurfsfassung eine bereits abgeschlossene Datenschutzprüfung darzustellen. Vor einer Veröffentlichung soll der Text den tatsächlich freigegebenen Ablauf beschreiben. Er darf nicht allein wegen einer knappen Projektfrist am Empfang ausgelegt werden.

Die Frage, ob für einzelne Verarbeitungsschritte eine Einwilligung eingeholt werden soll und was bei ihrer Ablehnung geschieht, ist in der bisherigen Abstimmung noch nicht beantwortet. Die Verwaltung bittet um eine gemeinsame Rückmeldung zu Verfahren, Information und Ansprechpartnern.

Aufgestellt von Johanna Klee | Kommunikation''')
add(14,'DSFA_Erhebung','docx','Erhebung für die Bewertung der Folgen des Vorhabens P01','Benedikt Fink und Nora Bergmann','01.10.2026','''1 Anlass und Umfang

Wir erfassen die geplante Verarbeitung von Sprachaufnahmen und Briefentwürfen durch Sprechfeder Clinical 2.4. Die Erhebung bezieht sich auf P01. Die KI in der Radiologie und die Versorgungsforschung werden in eigenen Vorgängen betrachtet. Der klinikweite Nutzen wird mit einer Entlastung bei der Dokumentation begründet. Ein gemessener Nutzen im eigenen Haus liegt noch nicht vor.

2 Betroffene und mögliche Auswirkungen

Betroffen wären behandelte Personen und Beschäftigte, deren Sprache, Nutzerkennung oder Bearbeitungszeiten verarbeitet werden. Fehlerhafte Zuordnung eines Entwurfs, nicht bemerkte Textveränderungen und ein zu weit reichender Supportzugriff wurden in der Aufnahme als mögliche Schäden beschrieben. Für Beschäftigte besteht die Frage, ob individuelle Bearbeitungszeiten zu einer Leistungsbewertung verwendet werden können. Die Funktionen der Auswertungsansicht hat der Anbieter noch nicht vollständig demonstriert.

3 Bestehende und vorgeschlagene Maßnahmen

Im angebotenen Aufbau bleibt die ärztliche Prüfung vor der Übernahme erforderlich. Die IT möchte automatische Übernahmen sperren, personenbezogene Nutzerzugänge einsetzen und Supportrechte zeitlich begrenzen. Ob diese Einstellungen im angebotenen Tarif verfügbar sind und exportierbar protokolliert werden, ist zu prüfen. Die Vorschläge sind noch nicht mit einem erfolgreichen Test oder einer Freigabe gleichzusetzen.

4 Offene Bewertungen

Für die Beurteilung von Erforderlichkeit und angemessenen Alternativen fehlen ein Vergleich mit der bestehenden Diktierlösung und die dokumentierte Betrachtung einer lokalen Verarbeitung. Die Speicherfristen aus dem Lieferantenvertrag sind noch nicht mit den klinischen Abläufen abgeglichen. Eintrittswahrscheinlichkeiten und verbleibende Risiken können wir nach der bloßen Vorführung nicht verlässlich beziffern. Insbesondere die Bedeutung von Fehlzuordnungen wird gemeinsam mit der ärztlichen Leitung zu prüfen sein.

5 Weitere Beteiligung

Herr Seidel liefert das Ergebnis der Zugriffsprüfung. Frau Dr. Feld beschreibt die Kontrollschritte in der Briefbearbeitung. Herr Lindner erhält die Beschäftigtenfunktionen. Herr Fink fasst nach Eingang dieser Beiträge seine Beratung zusammen. Dieses Erhebungsblatt dokumentiert den Beginn der Prüfung und enthält keine abgeschlossene Datenschutz-Folgenabschätzung.

Benedikt Fink''')
add(15,'Station_Rueckfragen','eml','Entlassbriefentwürfe und unsere Nachtbesetzung','Dr. Amira Feld <amira.feld@auenhoehe.example>','2026-10-01T09:26:00','''Liebe Nora,

für die Planung brauche ich noch eine klare Antwort, wer den Entwurf nachts prüfen soll. Wir haben dann nicht dieselbe Besetzung wie bei der Vorführung. Der Entwurf muss im System erkennbar unfertig bleiben, bis jemand ihn tatsächlich gelesen hat. Eine grüne Statusanzeige darf nicht den Eindruck erwecken, die KI habe den Inhalt bereits medizinisch geprüft.

Bitte messt später nicht nur die Zeit vom Diktat bis zum ersten Text. Wir müssen den Entwurf mit dem Ausgangsinhalt abgleichen und gegebenenfalls neu formulieren. Auch diese Zeit gehört dazu. Die zwanzig Briefe pro Woche wären eine Planmenge für eine spätere Messung, keine schon erreichte Zahl.

Ich kann mit zwei Kolleginnen selbst verfasste Beispieldiktate erstellen. Für den Beginn möchte ich weder alte Entlassbriefe kopieren noch echte Verläufe als vermeintlich anonyme Beispiele einstellen. Ob eine spätere Erprobung mit Patientendaten verantwortet werden kann, müssen wir vorher entscheiden.

Viele Grüße
Amira''')
add(16,'Einkauf_Vertrag','eml','Bestellung bleibt bis zur Rückmeldung gesperrt','Hagen Wolf <hagen.wolf@auenhoehe.example>','2026-10-01T13:42:00','''Hallo Nora,

ich habe für das Angebot über 10.140 EUR netto eine Vormerkung angelegt. Das ist keine Bestellung und keine Zahlungsfreigabe. Der Lieferant hat das telefonisch ebenfalls so verstanden. Die zwei Monate und dreißig Nutzer sind im Preis enthalten; die 4.800 EUR Einrichtung dürfen in der Kalkulation nicht noch einmal als Monatskosten auftauchen.

Die automatische Fortführung aus dem ersten Vertriebsblatt ist im aktuellen Angebot nicht mehr enthalten. Im Vertragsentwurf steht aber weiterhin die Verwendung zur Modellverbesserung. Solange wir keine abgestimmte Vertragsfassung haben, verschicke ich keine unterschriebene Annahme. Eine Zusicherung nur im Telefonat möchte ich nicht als Vertragsbestandteil ablegen.

Für P02 liegt mir noch kein vergleichbar vollständiges Angebot vor. Wenn ihr die beiden Projekte in einer Übersicht zeigt, bitte nicht den Preis von P01 einfach übertragen. Interne Personalstunden sind außerdem kein Rechnungsbetrag dieses Lieferanten.

Viele Grüße
Hagen''')
add(17,'Zweckbestimmung_Radiaviso','pdf','Zweckbestimmung Radialert Triage Version 5.0','Radiaviso Medical GmbH | Produktmanagement','18.09.2026','''Sehr geehrte Frau Dr. Feld,

Radialert Triage 5.0 unterstützt den ärztlichen Dienst bei der Priorisierung von Untersuchungen in einer Arbeitsliste. Die Anwendung verarbeitet die in der freigegebenen Schnittstellenspezifikation bezeichneten Bildserien und erzeugt eine zusätzliche Markierung. Die Markierung ersetzt weder die Befundung noch die Prüfung sämtlicher bereitgestellter Untersuchungen durch qualifiziertes Personal. Ohne Markierung darf eine Untersuchung nicht allein deshalb von der Bearbeitung ausgeschlossen werden.

Dieses Dokument beschreibt die Version 5.0. Die Softwarelieferung für Ihr Haus wird nach der gegenwärtigen Planung die Version 5.1 enthalten. Ein ergänzendes Dokument zum Unterschied zwischen diesen Versionen ist angefordert. Die hier beschriebene Zweckbestimmung darf deshalb nicht ohne Abgleich als abschließende Beschreibung der angelieferten Version verwendet werden. In der Vorführung war eine neu gestaltete Priorisierungsanzeige zu sehen.

Der Betrieb ist auf dem lokalen System des Hauses vorgesehen. Die Schnittstelle übergibt Untersuchung und interne Kennung gemeinsam. Wird die Verbindung zum Bildarchiv unterbrochen, muss die Arbeitsliste den Ausfall erkennbar anzeigen. Die Befundung muss über die vorhandenen klinischen Verfahren weiter möglich sein. Der Hersteller stellt Wartung und Updates nach einem gesonderten Plan bereit.

Die Unterlagen zur Konformität und die Gebrauchsanweisung werden als gesondertes Paket geliefert. Dieses Schreiben stellt keine Konformitätserklärung dar und enthält keine Bestätigung einer bereits erfolgten Abnahme im Krankenhaus. Bitte teilen Sie uns die konkret vorgesehenen Untersuchungsarten und Arbeitsabläufe mit, damit wir diese mit dem freigegebenen Einsatzbereich abgleichen können.

Mit freundlichen Grüßen
Dr. Eva Riedel | Produktmanagement''')
add(18,'Konformitaet_Lieferstand','docx','Begleitschreiben zum Stand der Produktunterlagen','Radiaviso Medical GmbH | Dr. Eva Riedel','28.09.2026','''Sehr geehrte Frau Bergmann,

wir bestätigen den Eingang Ihrer Bitte um das vollständige Produktdossier für Radialert Triage 5.1. Das Paket, das unser Vertrieb letzte Woche zusammenstellte, enthielt die Zweckbestimmung der Version 5.0 und eine Produktübersicht. Es war nicht als vollständige Übergabe aller für Ihren Einsatz benötigten Nachweise vorgesehen. Wir bitten darum, die Übermittlung nicht als Abschluss Ihrer Prüfung zu verbuchen.

1 Noch ausstehende Unterlagen

Die Rechtsabteilung stellt die geltende Konformitätserklärung und die zugehörige eindeutige Produktidentifikation zusammen. Die Beschreibung der Unterschiede zwischen Version 5.0 und 5.1 befindet sich beim Produktteam. Für die aktualisierte Arbeitslistenanzeige möchten wir Ihnen außerdem die gültige Gebrauchsanweisung übermitteln. Die Zuordnung etwaiger Bescheinigungen zum konkreten Produkt und ihrem Geltungsbereich wird in diesem Paket erläutert.

2 Einbindung in Ihre Umgebung

Der technische Ansprechpartner benötigt die Version des Bildarchivs und die Spezifikation der Übergabeschnittstelle. Die lokale Installation ist nicht identisch mit einer bereits geprüften Einbindung in Ihren klinischen Ablauf. Aus den von Ihnen beschriebenen Anforderungen an Kennungen und Priorisierungsanzeigen können sich zusätzliche Abstimmungen ergeben. Bitte schalten Sie die produktive Weitergabe von Bildern bis zur vereinbarten Installation nicht frei.

3 Weitere Aussage des Herstellers

Unsere Vertriebsangabe, die Anwendung werde als Medizinprodukt angeboten, ersetzt die angeforderten Unterlagen nicht. Eine eigene Erklärung Ihrer IT zur vollständigen regulatorischen Konformität erwarten wir nicht. Bitte prüfen Sie die Übergabe anhand der tatsächlichen Dokumente und teilen Sie uns Unstimmigkeiten mit. Wir beabsichtigen, den ergänzten Stand bis zum 8. Oktober bereitzustellen; diese Ankündigung ist noch keine erfolgte Lieferung.

Mit freundlichen Grüßen
Dr. Eva Riedel''')
add(19,'Radiaviso_Anhang','eml','Zweckbestimmung im Anhang und Versionsstand','Dr. Eva Riedel <eva.riedel@radiaviso.example>','2026-09-28T15:08:00','''Sehr geehrte Frau Bergmann,

ich übersende Ihnen noch einmal die Zweckbestimmung 5.0, damit der Bezug im Projektordner eindeutig bleibt. Das Dokument ist dieser Nachricht tatsächlich beigefügt. Der für Ihre Installation vorgesehene Stand heißt 5.1. Die Differenzbeschreibung ist noch nicht beigefügt; bitte betrachten Sie die beiden Nummern nicht als bloßen Schreibfehler.

Unsere technische Abteilung möchte am 8. Oktober mit Ihrer Medizintechnik die Schnittstelle besprechen. Dafür benötigen wir eine Übersicht der Felder und Versionen, jedoch keine Patientendaten. Die Übertragung von Bildkopien können wir zunächst mit eigens erstellten Testobjekten nachvollziehen.

Die Konformitätserklärung und die gültige Gebrauchsanweisung werden separat zusammengestellt. Ich möchte nicht, dass unsere Terminvereinbarung als Abnahme oder als Freigabe durch Ihr Haus verstanden wird. Senden Sie mir bitte vorab die Namen der verantwortlichen Personen für IT, Medizintechnik und ärztliche Leitung.

Freundliche Grüße
Eva Riedel''')
add(20,'Medizintechnik_Abgleich','docx','Abgleich der Produktlieferung mit dem geplanten Einsatz','Lena Auer | Medizintechnik','01.10.2026','''1 Vorliegender Stand

Wir haben die Produktbeschreibung Radialert Triage 5.0, eine E-Mail zur geplanten Version 5.1 und das Begleitschreiben zum Unterlagenstand erhalten. Eine vollständige Dokumentenübergabe wurde vom Hersteller ausdrücklich noch nicht bestätigt. Die im Projektordner enthaltene Vertriebsfolie genügt nicht für die Zuordnung der konkreten Installation.

2 Geplanter klinischer Ablauf

Die Anwendung soll eine Priorisierungsmarkierung in einer bestehenden Arbeitsliste ergänzen. Frau Dr. Feld verlangt, dass unmarkierte Untersuchungen weiterhin vollständig bearbeitet werden. Der Ausfall der Zusatzfunktion muss erkennbar sein. In der Demonstration blieb eine zuletzt angezeigte Markierung stehen, während der Dienst unterbrochen war. Ob dies nur an der Demonstrationsumgebung lag, soll im Installationstest geklärt werden.

3 Kennungen und Änderungen

Die Schnittstelle verwendet die Untersuchungskennung des Bildarchivs. Bei einem erneut gestarteten Auftrag darf eine alte Ausgabe nicht einer anderen Untersuchung zugeordnet werden. Die IT wird die Übergabe mit neu erzeugten Testobjekten untersuchen. Updates sollen erst nach dokumentierter Kenntnis der Änderungen und interner Abstimmung eingespielt werden. Ein automatischer Austausch des Modells ist in unserem derzeitigen Entwurf nicht vorgesehen.

4 Unterlagen für den nächsten Termin

Wir benötigen die aktuelle Zweckbestimmung, die Gebrauchsanweisung, die eindeutige Produktidentifikation und die zur gelieferten Version gehörenden regulatorischen Nachweise. Außerdem muss der Hersteller erklären, welche Anwenderschulung und welche örtlichen Voraussetzungen er verlangt. Die Prüfung der konkreten Einbindung und die Entscheidung über die Nutzung im Behandlungsablauf sind noch offen. Ein Beschaffungstermin ersetzt diese Schritte nicht.

5 Zuständigkeit

Die Medizintechnik dokumentiert den Eingang und den Versionsabgleich. Sie kann die medizinische Eignung für die vorgesehenen Abläufe nicht allein bestätigen. Die ärztliche Leitung, IT und Informationssicherheit werden ihre Beiträge gesondert zeichnen. Zum Stichtag dieses Vermerks besteht keine Freigabe für die produktive Verwendung.

Lena Auer''')
add(21,'Radiologie_Wunsch','eml','Priorisierung hilft uns nur bei erkennbaren Grenzen','Dr. Amira Feld <amira.feld@auenhoehe.example>','2026-10-02T08:33:00','''Hallo Nora,

wir möchten die zusätzliche Markierung in der Arbeitsliste grundsätzlich erproben. Ich möchte aber verhindern, dass aus der Reihenfolge eine stillschweigende Entscheidung über die Behandlung wird. Jede Untersuchung muss nach unserem üblichen Ablauf gelesen werden, auch wenn das Zusatzsystem nichts markiert.

Bei der Vorführung irritierte mich, dass die letzte Markierung nach der Verbindungsunterbrechung noch sichtbar blieb. Für unser Team muss erkennbar sein, ob die Anzeige aktuell ist. Wir brauchen dafür eine nachvollziehbare Beschreibung und einen Test mit der vorgesehenen Version. Die Unterlage 5.0 erklärt nicht ohne Weiteres, was die angekündigte Version 5.1 anders macht.

Bitte nehmt die Schulung und den Ausfallablauf mit in die Terminplanung auf. Eine kurze Einweisung im Vertriebsgespräch ist mir für den Dienstbetrieb zu wenig. Ich habe noch keine ärztliche Freigabe erteilt und möchte diese Entscheidung erst auf Grundlage der tatsächlichen Unterlagen treffen.

Viele Grüße
Amira''')
add(22,'Beinahevorfall','docx','Vermerk über eine verhinderte Übernahme in die Testumgebung','Nora Bergmann | IT','02.10.2026','''1 Beobachtung

Am 2. Oktober um 11:14 Uhr wollte ein Mitarbeiter für die Demonstration von P01 einen Text aus einem bereits freigegebenen Entlassbrief kopieren. Nora Bergmann bemerkte den geöffneten Dokumentausschnitt vor dem Einfügen in das Browserfenster. Sie bat den Mitarbeiter, den Vorgang abzubrechen. Nach seiner Erklärung wollte er die Unterschiede zur bestehenden Diktiersoftware an einem vertrauten Beispiel zeigen. Er hatte die Testregel nicht als Verbot dieser Verwendung verstanden.

2 Bisher gesicherter Ablauf

Der Text wurde nach gemeinsamer Durchsicht nicht in das Eingabefeld der Anwendung eingefügt. Die Browseransicht blieb leer. Die IT sicherte die vorhandenen Anwendungsprotokolle der Testkennung. Daraus ergibt sich für den fraglichen Zeitraum kein abgeschickter Diktatvorgang. Ob die Zwischenablage durch weitere lokal eingesetzte Funktionen verarbeitet wurde, wurde in diesem Termin nicht geprüft. Ein pauschaler Abschluss der technischen Prüfung erfolgte deshalb nicht.

3 Sofortige Schritte

Der Demonstrationstermin wurde unterbrochen. Für weitere Versuche wurden ausschließlich vorab neu geschriebene Beispielsätze in einem eigenen Ordner bereitgestellt. Herr Fink und Herr Seidel erhielten eine kurze Sachverhaltsmitteilung. Der Mitarbeiter wurde nicht aufgefordert, den Inhalt des echten Briefs zur Erklärung per E-Mail weiterzusenden. In diesem Vermerk werden keine Patientendaten festgehalten.

4 Weitere Klärung

Herr Seidel prüft bis zum 6. Oktober die für den Ablauf verfügbaren Protokolle und die lokale Zwischenablagekonfiguration. Herr Fink bewertet den dokumentierten Sachverhalt anschließend gesondert. Dieser Vermerk trifft noch keine abschließende Aussage zu Meldepflichten. Der Umstand, dass eine Übermittlung verhindert werden sollte, genügt für sich allein weder zur Feststellung noch zum Ausschluss jedes denkbaren Datenabflusses.

Nora Bergmann''')
add(23,'Ticket_Beinahevorfall','eml','Ticket IT 261002 44 und Protokollsicherung','Karim Seidel <karim.seidel@auenhoehe.example>','2026-10-02T12:06:00','''Hallo Nora,

ich habe für den heute unterbrochenen Versuch das Ticket IT-261002-44 eröffnet. Bitte lasse die Testkennung bis zum Abgleich unverändert. Die vorhandenen Protokolle zeigen keine abgeschickte Eingabe um 11:14 Uhr; das ist ein konkreter Befund zu diesem System und noch keine vollständige Prüfung sämtlicher lokaler Funktionen.

Den betroffenen Entlassbrief benötige ich nicht in der E-Mail. Wir dokumentieren den Ablauf mit Zeit, Kennung und den tatsächlich vorhandenen Protokollereignissen. Wenn wir für eine weitere Untersuchung Einsicht benötigen, vereinbaren wir dafür einen begrenzten Zugang. Bitte verschicke keinen Screenshot mit dem Briefinhalt in einen größeren Verteiler.

Ich prüfe die Zwischenablageeinstellungen und melde Herrn Fink den gesicherten Stand. Anschließend können wir sauber festhalten, welche Feststellungen vorliegen und welche nicht. Für weitere Demonstrationen sollten die selbst geschriebenen Beispielsätze bereits vor Öffnen der Anwendung bereitliegen.

Viele Grüße
Karim''')
add(24,'Forschungsprotokoll','docx','Projektbeschreibung SEPSIS LR in der Fassung vom 29 September','Dr. Oskar Wenck | Versorgungsforschung','29.09.2026','''1 Gegenstand der Untersuchung

Wir möchten retrospektiv untersuchen, welche organisatorischen Abläufe mit längeren stationären Aufenthalten bei den im Projekt definierten Fallgruppen zusammenhängen. Vorgesehen sind Daten aus den Jahren 2022 bis 2025. Die Auswertung soll Versorgungsabläufe beschreiben und die Planung weiterer wissenschaftlicher Untersuchungen unterstützen. Sie soll keine aktuellen Behandlungsentscheidungen auslösen und keine individuellen Therapieempfehlungen erzeugen.

2 Vorgesehener Datenbestand

Nach einer ersten Abfrage kommen voraussichtlich etwa 1.800 stationäre Fälle für eine weitere Prüfung infrage. Diese Zahl bezeichnet einen Suchbestand und noch keine freigegebene Studienpopulation. Der Datenkatalog nennt Alter in Gruppen, Zeitabstände, ausgewählte Befundkategorien und Entlassungsmerkmale. Namen, Kontaktdaten und Freitext sollen nicht an die Hochschule gegeben werden. Die für Rückfragen erforderliche Zuordnungstabelle soll getrennt im Krankenhaus bleiben.

3 Zusammenarbeit mit dem MVZ

Die spätere ambulante Versorgung könnte für unsere Fragestellung relevant sein. Das MVZ verfügt jedoch über einen eigenen Bestand. Umfang, Verknüpfung und Voraussetzungen einer Bereitstellung sind noch nicht vereinbart. Wir dürfen daher aus dem stationären Suchbestand noch keine zugesagte Zahl verknüpfter Verläufe ableiten. Die MVZ-Geschäftsführung möchte vor einem Export die konkrete Liste der Merkmale und die vorgesehene Nutzung erhalten.

4 Zusammenarbeit mit der Hochschule

Die Hochschule Saalebogen soll die statistische Auswertung begleiten. Wer die Forschungszwecke und Methoden verbindlich festlegt, ist im Kooperationsentwurf noch nicht abschließend geregelt. Eine eigene Verwendung für weitere Projekte wurde mündlich angesprochen. Das Training eines allgemein nutzbaren Modells wäre ein weiterer Verarbeitungsschritt, dessen Ziel und Empfänger wir bislang nicht ausreichend beschrieben haben.

5 Vorbereitungsstand

Die Ethikstelle hat bisher nur eine Anfrage erhalten. Ein Votum liegt nicht vor. Die rechtlichen Voraussetzungen und die erforderliche Information betroffener Personen werden parallel geprüft. Es wurde noch kein Datensatz an die Hochschule oder an einen kommerziellen Anbieter übermittelt. Der vorliegende Entwurf ist die Grundlage der weiteren Abstimmung.

Dr. Oskar Wenck''')
add(25,'Datenkatalog_Forschung','docx','Datenkatalog und Vorschlag zur Pseudonymisierung','Dr. Oskar Wenck und Nora Bergmann','30.09.2026','''1 Auswahl und Kennung

Der stationäre Suchbestand umfasst vorläufig rund 1.800 Fälle aus den Jahren 2022 bis 2025. Das Forschungsteam möchte jedem eingeschlossenen Fall eine zufällig erzeugte Studienkennung zuweisen. Die Zuordnung zwischen Kennung und internem Fallkennzeichen soll ausschließlich in einem getrennten, zugriffsbeschränkten Bereich des Krankenhauses liegen. Der Export soll keine Krankenhausfallnummer enthalten. Die technische Umsetzung ist noch nicht abgenommen.

2 Gewünschte Merkmale

Der erste Datenkatalog enthält Altersgruppen, Aufenthaltsdauer, Aufnahmeart, Zeitabstände zwischen ausgewählten dokumentierten Ereignissen und zusammengefasste Befundkategorien. Geschlecht wird als Merkmal angefragt; die wissenschaftliche Begründung der benötigten Ausprägungen steht noch aus. Exakte Aufnahme- und Entlassungstage sollen möglichst durch relative Zeitabstände ersetzt werden. Ob für eine saisonale Betrachtung Monat und Jahr benötigt werden, soll das Forschungsteam erläutern.

3 Grenzen des Entwurfs

Seltene Merkmalskombinationen können auch ohne Namen auffallen. Der aktuelle Katalog enthält noch keine festgelegte Mindestgröße für Auswertungsgruppen und keine verbindliche Regel für die Zusammenfassung seltener Kombinationen. Die Bezeichnung pseudonymisiert bedeutet hier, dass das Krankenhaus eine getrennte Zuordnung behält. Sie ist keine Behauptung, der Forschungsdatensatz sei unter allen Umständen anonym.

4 Ambulante Nachverfolgung

Für einen Abgleich mit dem MVZ fehlt ein vereinbartes Verfahren. Weder Versichertennummern noch Kontaktangaben sollen unbesehen in eine gemeinsame Datei kopiert werden. Vor einer Entscheidung benötigen wir eine Beschreibung, wer den Abgleich durchführt, welche Kennungen dabei entstehen und wann die dafür verwendeten Hilfsdaten wieder entfernt werden. Die eigenständige MVZ-Gesellschaft hat einem Export noch nicht zugestimmt.

5 Dokumentation der Bereitstellung

Vor jeder tatsächlichen Übergabe sollen Datenversion, Feldliste, Empfänger und Freigabe dokumentiert werden. Die vorliegende Datei enthält keine Datensätze einzelner Personen. Sie beschreibt ausschließlich den geplanten Aufbau.

Nora Bergmann''')
add(26,'Ethikanfrage','docx','Anfrage zur vorgesehenen retrospektiven Untersuchung','Dr. Oskar Wenck | Klinikverbund Auenhöhe GmbH','01.10.2026','''An die Ethikstelle der Hochschule Saalebogen

Sehr geehrte Damen und Herren,

wir bitten um Mitteilung des vorgesehenen Verfahrens für die Beratung und Prüfung unseres Projekts SEPSIS-LR. Die Projektbeschreibung und der Datenkatalog liegen als Arbeitsfassungen vor. Wir möchten die erforderlichen Unterlagen vor einer Datenbereitstellung vollständig zusammenstellen. Mit dieser Anfrage behaupten wir weder ein bereits vorliegendes Votum noch den Abschluss der datenschutzrechtlichen Prüfung.

1 Vorhaben

Geplant ist eine retrospektive Untersuchung von Versorgungsabläufen anhand stationärer Daten aus den Jahren 2022 bis 2025. Nach erster technischer Suche kommen ungefähr 1.800 Fälle für die weitere Prüfung infrage. Der Suchbestand ist noch keine endgültige Studienpopulation. Die Auswertung soll keine Behandlung einzelner Personen steuern. Eine mögliche Erweiterung um ambulante Verläufe aus einer rechtlich eigenständigen MVZ-Gesellschaft wird noch abgestimmt.

2 Geplante Zusammenarbeit

Die Hochschule soll die statistische Auswertung wissenschaftlich begleiten. Die Festlegung der Zwecke, der Methoden und der Entscheidungsrechte ist noch Gegenstand des Kooperationsentwurfs. Ein zusätzlicher Vorschlag zum Training eines allgemein nutzbaren Modells ist bislang nicht konkretisiert und soll nicht stillschweigend in die ursprüngliche Projektbeschreibung aufgenommen werden.

3 Erbetene Rückmeldung

Bitte teilen Sie uns mit, welche Fassungen der Projektbeschreibung, der Information betroffener Personen und der Kooperationsunterlagen Sie benötigen. Wir möchten außerdem wissen, wie spätere Änderungen des Datenumfangs oder der Fragestellung vorgelegt werden sollen. Die Prüfung der Rechtsgrundlagen und der tatsächlichen technischen Maßnahmen erfolgt im Haus gesondert und soll durch eine Beratung Ihrer Stelle nicht ersetzt werden.

Für Rückfragen stehe ich unter oskar.wenck@auenhoehe.example zur Verfügung. Eine Datenübermittlung an Ihre Stelle ist mit dieser Anfrage nicht verbunden.

Mit freundlichen Grüßen
Dr. Oskar Wenck''')
add(27,'Forschung_Modelltraining','eml','SEPSIS LR und die Idee für ein späteres Modell','Dr. Oskar Wenck <oskar.wenck@auenhoehe.example>','2026-10-02T14:38:00','''Liebe Nora,

die Kollegin der Hochschule hat gefragt, ob die Daten später auch zum Training eines allgemein verwendbaren Prognosemodells dienen könnten. Das wäre für sie wissenschaftlich interessant. In unserer derzeitigen Projektbeschreibung geht es dagegen um eine retrospektive Beschreibung der Versorgungsabläufe. Ich habe ihr deshalb noch keine Zusage gemacht.

Die etwa 1.800 Fälle sind bislang nur das Ergebnis einer ersten Suche. Nach Prüfung der Kriterien kann die Zahl kleiner werden. Für verknüpfte ambulante Verläufe haben wir noch gar keine verlässliche Zahl, weil das MVZ die Bereitstellung zunächst besprechen möchte.

Bitte behandelt das Modelltraining als zusätzlich zu beschreibende Idee. Ich werde Zweck, Merkmale, Empfänger und spätere Nutzung mit der Hochschule klären. Es wäre ungünstig, wenn aus dem Wort Forschung in der Überschrift schon eine Zustimmung zu jeder späteren Verwendung abgeleitet würde. Die Anfrage an die Ethikstelle habe ich versandt; eine Antwort oder ein Votum liegt noch nicht vor.

Viele Grüße
Oskar''')
add(28,'MVZ_Datenanfrage','eml','Ambulante Daten sind noch nicht zur Übergabe freigegeben','Ruth Demir <ruth.demir@auenhoehe-mvz.example>','2026-10-05T08:19:00','''Guten Morgen Nora,

ich habe Oskars Anfrage erhalten. Wir unterstützen das Vorhaben grundsätzlich gern, benötigen aber zunächst eine konkrete Beschreibung, welche Daten des MVZ gebraucht werden und wer anschließend Zugriff erhält. Unsere Behandlungen stehen in einem eigenen System. Die gemeinsame Muttergesellschaft führt nicht automatisch zu einer einzigen Patientenakte.

Bitte lasst niemanden mit einem allgemeinen IT-Zugang einen Export ziehen. Unsere Mitarbeitenden kennen den späteren Forschungszweck bislang nur aus der kurzen Besprechung. Ich möchte vor einer Freigabe den Datenkatalog, die vorgesehenen Empfänger und das Verfahren des Abgleichs sehen. Auch die Rückfragen von Patienten müssen wir beantworten können.

Eine vollständige Liste mit Namen und Versichertennummern werde ich nicht vorsorglich schicken. Wir sollten zuerst klären, welche Zuordnung tatsächlich notwendig ist. Ich kann am Mittwoch mit Benedikt und Oskar die offenen Punkte besprechen. Eine konkrete Exportfreigabe enthält diese Nachricht nicht.

Freundliche Grüße
Ruth Demir''')
add(29,'Kooperationsentwurf','docx','Kooperationsentwurf für die wissenschaftliche Auswertung','Klinikverbund Auenhöhe GmbH und Hochschule Saalebogen','02.10.2026','''1 Gegenstand

Die Klinikverbund Auenhöhe GmbH und die Hochschule Saalebogen beabsichtigen, im Projekt SEPSIS-LR die in der Projektbeschreibung bezeichneten Versorgungsabläufe zu untersuchen. Dieser Entwurf soll die Zusammenarbeit vorbereiten. Er ist nicht unterzeichnet. Die endgültige Bestimmung von Datenumfang, Verantwortlichkeiten und Verarbeitungsvoraussetzungen steht noch aus.

2 Beiträge der Beteiligten

Die Klinik soll die Auswahl der geeigneten Fälle anhand gemeinsam abgestimmter Kriterien vornehmen. Die Hochschule soll das Auswertungskonzept entwickeln und die statistische Bearbeitung unterstützen. Vor der Übergabe eines Datensatzes müssen beide Seiten schriftlich bestätigen, welche Fassung der Projektbeschreibung und des Datenkatalogs verbindlich ist. Die ambulanten Daten der Auenhöhe MVZ GmbH sind in diesem Entwurf noch nicht zugesagt.

3 Verwendung der Daten

Die Daten sollen für die vereinbarte Untersuchung verwendet werden. Die Hochschule hat ergänzend vorgeschlagen, gewonnene Daten und Ergebnisse für spätere Forschungsprojekte und die Entwicklung allgemein nutzbarer Modelle einzusetzen. Dieser Erweiterungsvorschlag ist noch nicht abgestimmt. Die Klinik hat eine Beschreibung der zusätzlichen Zwecke, der Empfänger und der vorgesehenen Dauer verlangt. Aus dem Entwurf darf keine Zustimmung zur Weitergabe an kommerzielle Partner abgeleitet werden.

4 Zugriffe und Veröffentlichung

Die Zuordnungstabelle soll im Krankenhaus verbleiben. Zugriffsberechtigte Personen auf Hochschulseite sind vor einer Bereitstellung namentlich zu benennen. Ergebnisse sollen nur in einer Form veröffentlicht werden, die vor der Freigabe auf die Offenlegung einzelner Fälle geprüft wurde. Welche Mindestgrößen und Aggregationsregeln hierfür gelten, ist im Auswertungskonzept noch festzulegen.

5 Ausstehende Vereinbarungen

Die Beteiligten haben noch nicht abschließend geregelt, wer Zwecke und wesentliche Mittel der einzelnen Verarbeitungsschritte festlegt und wer Auskünfte an betroffene Personen bearbeitet. Die Unterzeichnung soll erst nach Abstimmung dieser Rollen und der Lösch- und Rückgaberegeln erfolgen. Bis dahin ist keine Datenlieferung aus diesem Entwurf geschuldet.

Fassung zur gemeinsamen Besprechung am 7. Oktober 2026''')
add(30,'Betriebsrat_Anfrage','pdf','Auskunft zu Beschäftigtenfunktionen der geplanten Anwendungen','Paul Lindner | Betriebsrat des Klinikverbunds Auenhöhe','01.10.2026','''Sehr geehrte Frau Dr. Hartung, sehr geehrte Frau Bergmann,

der Betriebsrat bittet um vollständige Information über die Beschäftigtenfunktionen der vorgesehenen Anwendungen. Bei der Vorführung von Sprechfeder war eine Ansicht mit Bearbeitungszeiten, Anzahl der Diktate und Nutzerkennungen zu sehen. Uns ist nicht bekannt, ob diese Ansicht abschaltbar ist, wer Zugriff erhält und welche Auswertungen exportiert werden können. Eine Zusage, dass keine Leistungsbewertung beabsichtigt sei, beantwortet die Fragen zur technischen Funktion noch nicht.

Bitte übermitteln Sie uns die Rollenbeschreibung, die vorgesehenen Aufbewahrungsfristen der Nutzungsprotokolle und die geplanten Schulungsunterlagen. Wir möchten wissen, ob individuelle Korrekturquoten oder Arbeitszeiten sichtbar werden und ob diese Daten mit anderen Systemen zusammengeführt werden können. Die Beschäftigten müssen außerdem erkennen können, wann eine Tonaufnahme beginnt und endet.

Für die Radiologie benötigen wir eine Beschreibung, wie die Zusatzanzeige in den täglichen Ablauf eingebunden werden soll. Die Verantwortung und die Besetzung in Nacht- und Wochenenddiensten dürfen nicht allein durch die Bereitstellung einer Software ungeklärt bleiben. Schulungszeit ist in der Personalplanung zu berücksichtigen.

Wir schlagen eine gemeinsame Besprechung am 8. Oktober vor. Bis dahin bitten wir darum, weder die Einführung als beschlossen darzustellen noch Beschäftigte zur Verwendung eigener Endgeräte aufzufordern. Der Betriebsrat hat keine Vereinbarung zu einer produktiven Nutzung unterzeichnet. Wir sind an einer praktikablen Lösung interessiert, benötigen dafür aber die tatsächlichen Funktionen und die vorgesehenen Regeln.

Mit freundlichen Grüßen
Paul Lindner | Vorsitzender''')
add(31,'Betriebsrat_Termin','eml','Termin am 8 Oktober und Bildschirmansichten','Paul Lindner <paul.lindner@auenhoehe.example>','2026-10-05T10:04:00','''Hallo Nora,

der Termin am Donnerstag passt. Bitte bring die tatsächlichen Ansichten für die Nutzerverwaltung und die Statistik mit. Eine Präsentation mit allgemeinen Produktvorteilen hilft uns bei den offenen Fragen wenig. Wir möchten sehen, welche personenbezogenen Zahlen eine Stationsleitung abrufen könnte.

Wenn die Funktionen noch nicht konfiguriert sind, kennzeichne das bitte. Wir sollten zwischen einer möglichen Einstellung und einer bereits wirksam getesteten Begrenzung unterscheiden. Für die Schulung brauchen die Kolleginnen und Kollegen eingeplante Arbeitszeit; sie sollen das nicht zwischen zwei Aufgaben selbst ausprobieren müssen.

Ich habe bisher nur unser Auskunftsschreiben abgezeichnet, keine Zustimmung zur Einführung. Das heißt nicht, dass wir die Projekte verhindern möchten. Wir brauchen einen nachvollziehbaren Rahmen und eine Gelegenheit, die konkrete Fassung zu beraten. Für den Termin reichen selbst erstellte Beispiele ohne Patientendaten und ohne individuelle Beschäftigtenstatistiken aus dem echten Betrieb.

Viele Grüße
Paul''')
add(32,'Notfallprobe','docx','Protokoll der Ausfallprobe in der getrennten Testumgebung','Karim Seidel | Informationssicherheit','05.10.2026','''1 Aufbau

Am 5. Oktober wurde von 14:00 bis 14:45 Uhr die Rückkehr zur bisherigen Dokumentation besprochen und in der Testumgebung erprobt. Nora Bergmann, Karim Seidel und eine Mitarbeiterin des ärztlichen Dienstes nahmen teil. Es wurde ein neu geschriebener Beispieltext verwendet. Produktive Patientendaten und die reale Befundungsarbeitsliste waren nicht beteiligt.

2 Unterbrechung der Verbindung

Nach einer absichtlichen Unterbrechung zeigte der Browser einen Verbindungsfehler. Der zuletzt sichtbare Entwurf blieb im Fenster. Die Teilnehmer konnten daraus nicht sicher erkennen, ob noch eine Speicherung erfolgen würde. Der Text wurde nicht in ein produktives System übernommen. Ein erneutes Laden des Fensters entfernte die lokale Anzeige, ließ aber die Frage nach einem eventuell verbliebenen serverseitigen Entwurf offen.

3 Rückkehr zum vorhandenen Ablauf

Der ärztliche Dienst konnte die bisherige Diktiersoftware öffnen und einen neuen Testtext aufnehmen. Der Wechsel dauerte in dieser einzelnen Probe sechs Minuten einschließlich Rückfrage bei der IT. Diese Beobachtung ist keine repräsentative Zeitmessung und wird nicht als Ergebnis des geplanten achtwöchigen Piloten verbucht. Die tatsächliche Nachtbesetzung war an dem Termin nicht beteiligt.

4 Offene Punkte

Die Kennzeichnung veralteter Entwürfe, die Vermeidung doppelter Übernahmen und die Bearbeitung unklar gespeicherter Inhalte müssen im Ablauf beschrieben werden. Für die Radiologie ist eine eigene Ausfallprobe mit der vorgesehenen Produktversion erforderlich. Die heute gezeigte Rückkehr zur bisherigen Diktiersoftware bestätigt nicht den Ausfallablauf der radiologischen Priorisierung.

5 Fortsetzung

Die IT fordert vom Anbieter eine Beschreibung der Speichervorgänge bei Verbindungsabbruch an. Frau Dr. Feld schlägt die Teilnehmer für einen späteren Test unter realistischen Dienstbedingungen vor. Bis diese Schritte und die übrigen Prüfungen abgeschlossen sind, bleibt die technische Probe eine Vorbereitung ohne Betriebsfreigabe.

Karim Seidel''')
add(33,'ISB_Rueckmeldung','eml','Zugänge, Ausfallprobe und nächste technische Schritte','Karim Seidel <karim.seidel@auenhoehe.example>','2026-10-06T08:05:00','''Guten Morgen Nora,

die Ausfallprobe ist dokumentiert. Wir konnten zur vorhandenen Diktiersoftware zurückkehren, haben aber noch offene Fragen zu dem im Browser sichtbaren Restentwurf. Die sechs Minuten aus diesem einzelnen Versuch dürfen bitte nicht als durchschnittliche Ausfallzeit oder gesicherte Leistungskennzahl in die Pilotwertung eingehen.

Für P02 fehlt uns weiterhin der Abgleich zur Version 5.1. Die Sitzung zur Schnittstelle können wir vorbereiten; dafür reichen technische Feldnamen und eigens erzeugte Testobjekte. Eine Öffnung des allgemeinen Fernwartungszugangs für den Hersteller ist noch nicht erforderlich und von mir nicht freigegeben.

Zum verhinderten Einfügen vom Freitag habe ich den bisherigen Protokollbefund an Benedikt weitergegeben. Die Prüfung der lokalen Zwischenablagefunktion ist noch nicht vollständig abgeschlossen. Bitte führe diesen Punkt als offen weiter, damit wir nicht aus einer leeren Anwendungsprotokollzeile einen umfassenden Abschluss ableiten.

Viele Grüße
Karim''')
add(34,'Rechtsabteilung_Rollen','eml','Drei Vorhaben benötigen drei nachvollziehbare Entscheidungen','Friederike Blum <friederike.blum@auenhoehe.example>','2026-10-06T09:12:00','''Hallo Nora,

ich habe die bisherigen Unterlagen gelesen. Für die Entscheidungsvorlage sollten wir die drei Vorhaben getrennt halten. Der Cloudvertrag für Briefentwürfe beantwortet weder die Fragen zur Produktversion in der Radiologie noch die Rollen der Hochschule bei der Forschung. Auch eine vom Anbieter übersandte Vereinbarung mit der Überschrift Auftragsverarbeitung klärt nicht jede tatsächlich vorgesehene Verwendung.

Bitte sammele die konkreten Fragen zur Rechtsgrundlage und zur Information der betroffenen Personen mit dem jeweils zugehörigen Datenfluss. Für die Forschung brauche ich besonders die Antwort, wer Zwecke und Auswertungsmethoden festlegt und ob das spätere Modelltraining überhaupt Teil des jetzigen Projekts sein soll.

Für Freitag kann ich einen abgestimmten Sachstand zuliefern. Einen fertigen positiven Konformitätsvermerk kann ich aus den angekündigten, aber noch nicht vorliegenden Nachweisen nicht erstellen. Wenn die Geschäftsführung einen eingeschränkten nächsten Vorbereitungsschritt beschließt, muss dessen Umfang erkennbar sein.

Viele Grüße
Friederike''')
add(35,'IT_Aktenuebergabe','eml','Unterlagen für die Lenkungsrunde am Freitag','Nora Bergmann <nora.bergmann@auenhoehe.example>','2026-10-06T13:24:00','''Guten Tag zusammen,

ich habe den Stand der drei Vorhaben zum heutigen Nachmittag zusammengeführt. Die vier ergänzenden Arbeitsmappen trennen Vorhaben und Datenflüsse, Maßnahmen und Freigaben, Lieferantennachweise sowie Pilotmessung und Kosten. Offene Angaben bleiben offen; ich habe keine angekündigte Lieferung als bereits eingegangen markiert.

Die Kostenplanung für P01 enthält dreißig Nutzer zu 89 EUR netto im Monat für zwei Monate, 4.800 EUR Einrichtung und eine interne Planannahme von 48 Stunden zu 65 EUR. Die internen Stunden sind keine Lieferantenrechnung. Die Messung mit zwanzig Zielbriefen je Woche ist nur geplant; alle acht Wochen haben noch keine beobachteten Messwerte.

Bitte ergänzt eure Beiträge mit Datum und dem zugehörigen Dokument. Ich benötige insbesondere die tatsächlichen Nachweise des Anbieters, den Versionsabgleich der Radiologie und die Beschreibung der Forschungserweiterung. Für den 19. Oktober habe ich keinen Produktivzugang und keine Patientendatenübermittlung veranlasst.

Freundliche Grüße
Nora Bergmann''',to='Dr. Selma Hartung <selma.hartung@auenhoehe.example>')
add(36,'Lenkungsrunde_Stand','docx','Vorbereitung der Lenkungsrunde am 9 Oktober','Nora Bergmann | Gesamtleitung IT','06.10.2026','''1 Auftrag für die Besprechung

Die Geschäftsführung soll über die nächsten Schritte der drei Vorhaben entscheiden. Bis zum Stichtag 6. Oktober 2026 liegt für keines der Vorhaben eine Betriebsfreigabe vor. Die Vorbereitung wurde mit eigens geschriebenen Texten und technischen Testobjekten durchgeführt. In den ergänzenden Arbeitsmappen werden Entscheidungen erst nach einer dokumentierten Erklärung eingetragen.

2 P01 Kosten und Messplan

Das Angebot nennt dreißig Nutzer zu 89,00 EUR netto pro Monat für zwei Monate sowie einmalig 4.800,00 EUR netto für die Einrichtung. Die externen Nettokosten betragen damit 10.140,00 EUR. Für die Planung setze ich zusätzlich 48 interne Stunden zu einem kalkulatorischen Satz von 65,00 EUR an. Das ergibt 3.120,00 EUR internen Aufwand und zusammen 13.260,00 EUR kalkulatorische Nettokosten. Diese interne Bewertung ist kein zusätzlicher Lieferantenanspruch.

Nach einer Freigabe sollen acht Wochen beobachtet werden. Als Planmenge sind zwanzig Briefe je Woche vorgesehen. Gemessen werden sollen die bisherige Bearbeitungszeit und die Zeit mit Unterstützung einschließlich Prüfung und Korrektur. Die Arbeitsmappe enthält noch keine Beobachtungswerte. Eine angenommene Zeitersparnis darf deshalb heute weder als Einsparung noch als Personalabbau eingeplant werden. Ein entlastender Effekt müsste außerdem im Dienst tatsächlich nutzbar sein.

3 Ausstehende Beiträge

Für P01 fehlen unter anderem der abgestimmte Vertrag zur Modellverwendung, die vollständigen Angaben zu Supportzugriffen und die Zuordnung des Plattformprüfberichts. Bei P02 stehen der Versionsabgleich und die Produktunterlagen aus. Bei P03 müssen Datenumfang, ambulante Verknüpfung, Rollen und zusätzliche Modellnutzung beschrieben werden. Die Ethikanfrage ist noch ohne Antwort.

4 Besprechungsvorbereitung

Jeder Fachbereich soll seinen Beitrag am Freitag erläutern und die dafür verwendeten Unterlagen benennen. Der Wunschstart am 19. Oktober bleibt ein Plantermin. Wenn nur ein begrenzter Vorbereitungsschritt beschlossen wird, soll der Beschluss den erlaubten Umfang, die verantwortliche Person und die noch ausstehenden Voraussetzungen ausdrücklich nennen. Ein vollständiger Betrieb aller drei Vorhaben lässt sich aus der Projektanlage allein nicht ableiten.

Nora Bergmann''')

for d in DOCS:
    if d['n'] == 11: d['to'] = 'Maja Roth | Sprechfeder Health Europe GmbH'
    if d['n'] == 26: d['to'] = 'Ethikstelle der Hochschule Saalebogen'

# A4 und Times New Roman entsprechen dem verbindlichen Repositoryformat.
def docx_file(d):
    doc=Document(); s=doc.sections[0]
    s.page_width=Cm(21); s.page_height=Cm(29.7)
    s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
    for sn in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[sn];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
        for script in ('ascii','hAnsi','eastAsia','cs'):st.element.get_or_add_rPr().rFonts.set(qn('w:'+script),'Times New Roman')
        for key in list(st.element.get_or_add_rPr().rFonts.attrib):
            if key.endswith('Theme'): del st.element.get_or_add_rPr().rFonts.attrib[key]
        st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.04
    doc.styles['Title'].font.size=Pt(15);doc.styles['Title'].font.bold=True
    for b in doc.styles.element.xpath('.//w:pBdr'):b.getparent().remove(b)
    doc.add_paragraph(d['author']);doc.add_paragraph(d['to']);doc.add_paragraph(d['date'])
    doc.add_paragraph(d['title'],'Title')
    for block in d['text'].split('\n\n'):
        heading=bool(re.match(r'^\d+ ',block)) and '\n' not in block and len(block)<90
        p=doc.add_paragraph(block,'Heading 1' if heading else 'Normal');p.paragraph_format.widow_control=True
        if heading:p.paragraph_format.space_before=Pt(8)
    doc.core_properties.author=doc.core_properties.last_modified_by='Klotzkette'
    doc.core_properties.title=d['title'];doc.core_properties.language='de-DE';doc.core_properties.created=doc.core_properties.modified=STAMP
    doc.save(CASE/d['name'])
def pdf_file(d):
    body=ParagraphStyle('body',fontName='TNR',fontSize=11,leading=14,spaceAfter=8)
    title=ParagraphStyle('title',parent=body,fontName='TNRB',fontSize=15,leading=18,spaceAfter=13)
    flow=[Paragraph(escape(d['author']),body),Paragraph(escape(d['to']),body),Paragraph(d['date'],body),Spacer(1,6),Paragraph(escape(d['title']),title)]
    flow += [Paragraph(escape(b).replace('\n','<br/>'),body) for b in d['text'].split('\n\n')]
    def footer(c,doc):c.setFont('TNR',9);c.drawRightString(A4[0]-56,30,str(doc.page))
    SimpleDocTemplate(str(CASE/d['name']),pagesize=A4,leftMargin=56,rightMargin=56,topMargin=48,bottomMargin=48,title=d['title'],author='Klotzkette').build(flow,onFirstPage=footer,onLaterPages=footer)
def mail_file(d):
    m=EmailMessage(policy=SMTP);m['From']=d['author'];m['To']=d['to'] if '<' in d['to'] else 'Nora Bergmann <nora.bergmann@auenhoehe.example>'
    m['Date']=format_datetime(datetime.fromisoformat(d['date']+'+02:00'));m['Subject']=d['title'];m['Message-ID']=f'<auenhoehe-{d["n"]:02d}-20261006@auenhoehe.example>';m['Content-Language']='de-DE'
    m.set_content(d['text'],charset='utf-8')
    attach={8:'04_Angebot_Sprechfeder.pdf',19:'17_Zweckbestimmung_Radiaviso.pdf'}
    if d['n'] in attach:
        p=CASE/attach[d['n']];m.add_attachment(p.read_bytes(),maintype='application',subtype='pdf',filename=p.name)
    (CASE/d['name']).write_bytes(m.as_bytes())
def main():
    pdfmetrics.registerFont(TTFont('TNR',str(serif_font_path())));pdfmetrics.registerFont(TTFont('TNRB',str(serif_font_path(bold=True))))
    for d in DOCS:
        {'docx':docx_file,'pdf':pdf_file,'eml':mail_file}[d['kind']](d)
    qa=ROOT/'quality/krankenhaus-it-ki';qa.mkdir(parents=True,exist_ok=True)
    manifest=[{k:d[k] for k in ['n','name','kind','title','author','date']}|{'sha256':hashlib.sha256((CASE/d['name']).read_bytes()).hexdigest(),'characters':len(d['text'])} for d in DOCS]
    (qa/'testakte-originale.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'files':len(DOCS),'types':{k:sum(d['kind']==k for d in DOCS) for k in ['docx','pdf','eml']},'case':str(CASE)},ensure_ascii=False))
if __name__=='__main__':main()
