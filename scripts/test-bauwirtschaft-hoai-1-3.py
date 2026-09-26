#!/usr/bin/env python3
"""Offline-Regression HOAI 1 bis 3; optionale Exporte und Office-Proben. Autor: Klotzkette."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
from email import policy
from email.parser import BytesParser
import io
from pathlib import Path
import re
import shutil
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageStat
from pypdf import PdfReader
import yaml
import zipfile

from readme_decimal_headings import normalize_decimal_headings
from testakte_disclaimer import NOTICE_BYTES, NOTICE_MARKDOWN, pdf_content_errors
from testakte_file_filter import include_in_working_dump
import importlib.util

HERE = Path(__file__).resolve().parent


def module(name, filename):
    spec = importlib.util.spec_from_file_location(name, HERE / filename)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


P = module("hoai123_pakete_test", "build-bauwirtschaft-hoai-1-3-pakete.py")
B = P.B
SKILLS = {
    1: "hoai-1-grundlagen-und-planungsauftrag-klaeren",
    2: "hoai-2-vorplanung-und-varianten-entwickeln",
    3: "hoai-3-entwurf-und-kostenberechnung-abstimmen",
}
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--assets", type=Path, help="Verzeichnis kanonischer Release-Exporte")
parser.add_argument("--native", action="store_true", help="Office-Neuberechnung in eigenen Wegwerfkopien")
args, remaining = parser.parse_known_args()


def contents(path):
    if path.suffix == ".pdf":
        return " ".join(" ".join(p.extract_text() or "" for p in PdfReader(path).pages).split())
    if path.suffix == ".docx":
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs] +
                         [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if path.suffix == ".eml":
        msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        return msg.get_body(preferencelist=("plain",)).get_content()
    return path.read_text(encoding="utf-8-sig")


class Originals(unittest.TestCase):
    def test_inventory_and_ascii_filenames(self):
        for phase, case in B.CASES.items():
            folder = B.ROOT / "testakten" / case["slug"]
            actual = {p.name for p in folder.rglob("*")
                      if include_in_working_dump(p, folder, include_gesamt_pdf=False)}
            self.assertEqual(actual, set(case["files"]))
            self.assertEqual(len(actual), 16)
            self.assertEqual({Path(n).suffix for n in actual},
                             {".pdf", ".docx", ".eml", ".csv", ".txt", ".png", ".xlsx"})
            for name in actual:
                self.assertTrue(name.isascii(), name)
                self.assertGreater((folder / name).stat().st_size, 250, name)

    def test_skill_scope_and_standalone_prompts(self):
        for phase, slug in SKILLS.items():
            path = B.ROOT / "bauwirtschaft/skills" / slug / "SKILL.md"
            text = path.read_text(encoding="utf-8")
            meta = yaml.safe_load(text.split("---", 2)[1])
            self.assertEqual(set(meta), {"name", "description"})
            self.assertEqual(meta["name"], slug)
            self.assertTrue(80 <= len(meta["description"]) <= 1024)
            self.assertLess(len(text.splitlines()), 500)
            self.assertEqual(len(re.findall(r"^## [1-6]\. ", text, re.M)), 6)
            self.assertNotIn("werkstatt.md", text)
            self.assertNotIn("Sourcecards", text)
            for link in re.findall(r"\]\(([^)]+)\)", text):
                if not link.startswith("https://"):
                    self.assertTrue((path.parent / link).resolve().is_file(), link)
            prompt = (B.ROOT / f"bauwirtschaft/bauwirtschaft-hoai-{phase}-werkstatt.md").read_text()
            self.assertEqual(len(re.findall(r"^# ", prompt, re.M)), 1)
            self.assertTrue(20000 <= len(prompt.encode()) <= 32000)
            for link in re.findall(r"\]\(([^)]+)\)", prompt):
                self.assertTrue(link.startswith("https://"), link)
            for required in ["Innenräume", "Gebäude", "Paragraf 34", "650p", "Anlage 10", "Rückfragen"]:
                self.assertIn(required, prompt)
            self.assertNotIn(chr(167), text + prompt)
            source = (B.ROOT / f"bauwirtschaft/references/hoai-{phase}-fachquellen.md").read_text()
            self.assertIn("gesetze-im-internet.de", source)
            self.assertIn("25.09.2026", source)

    def test_readmes_and_rubrics(self):
        injector = module("hoai123_inject_test", "inject-gesamt-pdf-section.py")
        for phase, case in B.CASES.items():
            folder = B.ROOT / "testakten" / case["slug"]
            text = (folder / "README.md").read_text()
            self.assertIn("<!-- decimal-headings -->", text)
            self.assertIn("<!-- reserved-example-contacts -->", text)
            self.assertEqual(text, normalize_decimal_headings(text))
            title = re.search(r"^# (.+)$", text, re.M).group(1)
            self.assertIn(f"HOAI Leistungsphase {phase}", title)
            self.assertIn(case["city"], title)
            self.assertIn("erfunden", text)
            self.assertIn(NOTICE_MARKDOWN + "\n\n| Was", text)
            self.assertEqual(text.count("<!-- BEGIN gesamt-pdf-section (autogen) -->"), 1)
            self.assertEqual(text.count("| Was | Format | Quelle |"), 1)
            for name in case["files"]:
                self.assertIn(f"[{name}]({name})", text)
            # Testet den Injector ausschließlich an einer eigenen temporären Kopie.
            with tempfile.TemporaryDirectory(prefix="hoai123-readme-") as tmp:
                target = Path(tmp)
                readme = target / "README.md"
                readme.write_text(text, encoding="utf-8")
                (target / "01_Original.pdf").touch()
                if (folder / "gesamt-pdf" / (case["slug"] + "_gesamt.pdf")).exists():
                    (target / "gesamt-pdf").mkdir()
                    (target / "gesamt-pdf" / (case["slug"] + "_gesamt.pdf")).touch()
                self.assertEqual(injector.inject(readme, case["slug"]), "unchanged")
            rubric = yaml.safe_load((folder / "rubric.yaml").read_text())
            self.assertEqual(rubric["plugin"], "bauwirtschaft")
            ids = [c["id"] for c in rubric["checks"]]
            self.assertEqual(len(ids), len(set(ids)))
            self.assertTrue(any(c["check_type"] == "human_review" for c in rubric["checks"]))
            for check in rubric["checks"]:
                self.assertIn(check["check_type"], {"working_file_count", "file_exists", "human_review"})
                if check["check_type"] == "file_exists":
                    self.assertTrue((folder / check["path"]).is_file())

    def test_documents_and_metadata(self):
        for phase, case in B.CASES.items():
            for number, name in enumerate(case["files"], 1):
                path = P.original(phase, number)
                if path.suffix not in {".pdf", ".docx", ".eml", ".txt"}:
                    continue
                text = contents(path)
                self.assertNotIn(chr(167), text, name)
                self.assertNotIn("Diese Testakte wurde", text, name)
                if path.suffix in {".pdf", ".docx"}:
                    self.assertIn("gez.", text, name)
                if path.suffix == ".pdf":
                    self.assertEqual(PdfReader(path).metadata.author, "Klotzkette")
                    self.assertFalse(pdf_content_errors(path.read_bytes()))
                    self.assertGreater(len(text), 250, name)
                elif path.suffix == ".docx":
                    doc = Document(path)
                    self.assertEqual(doc.core_properties.author, "Klotzkette")
                    self.assertEqual(doc.styles["Normal"].font.name, "Times New Roman")
                    self.assertEqual(doc.styles["Normal"].font.size.pt, 11)
                    self.assertGreater(len(text), 1300, name)
                elif path.suffix == ".eml":
                    msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                    for key in ["From", "To", "Date", "Subject", "Message-ID"]:
                        self.assertTrue(msg[key], (name, key))
                    self.assertIn(".example", msg["From"])
                    self.assertIn(".example", msg["To"])
                    self.assertIn("Mit freundlichen Grüßen", text)
                    self.assertGreater(len(text), 650, name)

    def test_images_and_csv(self):
        for phase, case in B.CASES.items():
            for number, name in enumerate(case["files"], 1):
                path = P.original(phase, number)
                if path.suffix == ".png":
                    with Image.open(path) as im:
                        self.assertGreaterEqual(im.width, 2000)
                        self.assertGreaterEqual(im.height, 1400)
                        self.assertEqual(im.info.get("Author"), "Klotzkette")
                        self.assertGreater(min(ImageStat.Stat(im.convert("RGB")).stddev), 20)
                elif path.suffix == ".csv":
                    with path.open(encoding="utf-8-sig", newline="") as stream:
                        rows = list(csv.reader(stream, delimiter=";"))
                    self.assertGreater(len(rows), 5)
                    self.assertTrue(all(len(r) == len(rows[0]) for r in rows))

    def test_cached_formulas_and_independent_arithmetic(self):
        for (phase, number), spec in P.WORKBOOKS.items():
            path = P.original(phase, number)
            formula = load_workbook(path, data_only=False)
            values = load_workbook(path, data_only=True)
            s, v = formula[spec["sheet"]], values[spec["sheet"]]
            self.assertEqual(formula.properties.creator, "Klotzkette")
            self.assertEqual(formula.properties.lastModifiedBy, "Klotzkette")
            self.assertEqual(s[spec["output"]].data_type, "f")
            self.assertAlmostEqual(v[spec["output"]].value, spec["baseline"], places=2)
            self.assertGreater(len(P.formulas(path)), 10)
            self.assertFalse(s.sheet_view.showGridLines)
            for row in v:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "e", (path.name, cell.coordinate))
            if spec["sheet"] == "Termine":
                self.assertEqual(v["D13"].value, datetime(2027, 10, 7))
                self.assertEqual(sum(v[f"C{r}"].value for r in range(6, 14)), 479)
                self.assertIn('NA()', s["D6"].value)
            else:
                end = 18 if phase == 3 else 14
                for row in range(6, end):
                    self.assertAlmostEqual(v[f"E{row}"].value,
                                           v[f"B{row}"].value * v[f"D{row}"].value, places=2)
                self.assertIn("407", s[spec["output"]].number_format)
                if phase == 1:
                    self.assertAlmostEqual(v["E20"].value, 3931.8, places=2)
                if phase == 2:
                    self.assertAlmostEqual(v["E19"].value, 3167645.8, places=2)
                    self.assertAlmostEqual(v["G19"].value + 48000 * 1.19, 3190351, places=2)
                if phase == 3:
                    self.assertEqual(v["E26"].value, 26800)
                    self.assertEqual(v["E28"].value, 259908)
            formula.close()
            values.close()

    def test_project_specific_evidence_and_chronology(self):
        self.assertIn("120", contents(P.original(1, 1)))
        self.assertIn("160", contents(P.original(1, 14)))
        self.assertIn("Dachboden", contents(P.original(1, 15)))
        self.assertIn("F. Linde", contents(P.original(1, 6)))
        self.assertNotIn("Beitrag vom 16.03.", contents(P.original(2, 6)))
        self.assertIn("75", contents(P.original(2, 2)))
        self.assertIn("verbleiben", contents(P.original(2, 14)))
        self.assertIn("kein Vorbescheid", contents(P.original(2, 9)))
        self.assertNotIn("Raumliste vom 05.06.", contents(P.original(3, 6)))
        self.assertNotIn("Beitrag L-04 vom 03.06.", contents(P.original(3, 6)))
        self.assertIn("12,20 bis 14,00", contents(P.original(3, 8)))
        self.assertIn("1,80 mal 2,00", contents(P.original(3, 9)))
        self.assertIn("keine Zustimmung", contents(P.original(3, 13)))
        self.assertIn("kein Auftrag", contents(P.original(3, 15)))

    @unittest.skipUnless(args.assets, "Exportprüfung mit --assets DIR")
    def test_canonical_exports(self):
        for phase, case in B.CASES.items():
            slug = case["slug"]
            total_name = slug + "_gesamt.pdf"
            total = B.ROOT / "testakten" / slug / "gesamt-pdf" / total_name
            self.assertFalse(pdf_content_errors(total.read_bytes()))
            for suffix in ["", "-einzelpdfs"]:
                archive = args.assets / f"testakte-{slug}{suffix}.zip"
                with zipfile.ZipFile(archive) as z:
                    names = z.namelist()
                    expected = ({Path(n).stem + ".pdf" for n in case["files"]} if suffix
                                else set(case["files"]) | {total_name}) | {"README.txt"}
                    self.assertEqual(set(names), expected)
                    self.assertEqual(len(names), len(expected))
                    self.assertTrue(all("/" not in n for n in names))
                    self.assertEqual(z.read("README.txt"), NOTICE_BYTES)
                    for number, name in enumerate(case["files"], 1):
                        if suffix:
                            self.assertFalse(pdf_content_errors(z.read(Path(name).stem + ".pdf")))
                        else:
                            self.assertEqual(z.read(name), P.original(phase, number).read_bytes())
                    if not suffix:
                        self.assertEqual(z.read(total_name), total.read_bytes())

    @unittest.skipUnless(args.native, "Office-Mutationsprüfung mit --native")
    def test_native_input_mutations_zero_and_blank(self):
        zero_expected = {(1, 12): 7032.9, (2, 10): 3133231, (3, 11): 2116800, (3, 14): -1}
        with tempfile.TemporaryDirectory(prefix="hoai123-native-test-") as tmp:
            directory = Path(tmp)
            for mode in ["changed", "zero", "blank"]:
                stage = directory / mode
                stage.mkdir()
                probes = []
                specs = []
                for key, spec in P.WORKBOOKS.items():
                    source = P.original(*key)
                    target = stage / source.name
                    shutil.copyfile(source, target)
                    mutation = spec["changed"] if mode == "changed" else 0 if mode == "zero" else "blank"
                    P.patch_xlsx(target, spec, mutation=mutation)
                    probes.append(target)
                    specs.append((key, spec, P.formulas(source)))
                for output, (key, spec, formulas) in zip(P.roundtrip(probes, stage), specs):
                    self.assertEqual(P.formulas(output), formulas)
                    book = load_workbook(output, data_only=True)
                    value = book[spec["sheet"]][spec["output"]].value
                    if mode == "blank":
                        self.assertEqual(value, "#N/A", (key, value))
                    else:
                        expected = spec["expected"] if mode == "changed" else zero_expected[key]
                        self.assertAlmostEqual(value, expected, places=2, msg=str((key, mode)))
                        self.assertFalse([c.coordinate for row in book[spec["sheet"]]
                                          for c in row if c.data_type == "e"])
                    book.close()


if __name__ == "__main__":
    unittest.main(argv=[__file__, *remaining], verbosity=2)
