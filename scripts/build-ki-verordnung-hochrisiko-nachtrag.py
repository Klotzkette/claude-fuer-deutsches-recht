#!/usr/bin/env python3
"""Zwei neue Hochrisiko-Akten. Alte Recruitingakte bleibt unangetastet.

Erst ohne --messages, dann MJS-Tabellenbuilder, danach --messages ausführen.
Keine Musterlösung, keine Gesamt-PDF und keine Archive im Nutzercorpus.
"""
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email.headerregistry import Address
from email.policy import SMTP
from email.utils import getaddresses
import argparse, hashlib, json, re, html, os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

ROOT=Path(__file__).resolve().parents[1]
WARNING='> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n>\n> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'
TAG='ki-verordnung-v445.35.0'

def record(file,title,author,body):
    return dict(file=file,title=title,author=author,body=body.strip())
def email(file,title,sender,to,date,body,attachments=()):
    return dict(file=file,title=title,sender=sender,to=to,date=date,body=body.strip(),attachments=list(attachments))

CASES=[dict(
slug='ki-hochrisiko-justizassistenz-jena',title='Jena Aktenklar mit neuer Entscheidungsassistenz',reference='JA-2026-104',
intro='Eine fiktive Justizprojektstelle in Jena prüft mehrere Versionen von Aktenklar. Aus Paginierung und Registerbildung wird eine Unterstützung für Entscheidungsentwürfe. Der Prüfauftrag verlangt eine begründete Einordnung und konkrete Anbieterfragen; zwischen Vertriebsangaben und Testprotokollen bestehen Unterschiede.',
docs=[
record('02_Leistungsbeschreibung.docx','Aktenklar Leistungsumfang der Versionen','Friedemann Sturm | Aktenklar Software GmbH | Lichtbogenweg 8 | 07743 Jena','''An die Projektstelle Justizdigitalisierung Jena
Vertragsanlage zum Angebot AK-26-114, Stand 18. September 2026

## 1 Vereinbarter Grundumfang

Aktenklar 1.4 erkennt mittels eines trainierten Layoutmodells Dokumentgrenzen in eingescannten Akten, liest Überschriften aus und erzeugt ein Inhaltsverzeichnis. Die Originaldateien werden unverändert im Lesefenster daneben geöffnet. Die Ausgabe nennt Dokumentbezeichnung, Blattzahl und den Fundort. Eine Bewertung des Vorbringens oder der beteiligten Personen ist für diese Version nicht vorgesehen. Die Geschäftsstelle kann falsche Blattzahlen manuell berichtigen; das Ursprungsregister bleibt im Änderungsprotokoll erhalten.

Die Projektstelle darf die Version mit synthetischen Akten in ihrem abgetrennten Testbereich erproben. Es sind weder eine Gerichtsentscheidung noch eine automatische Kommunikation mit Parteien beauftragt. Die Laufzeit vom 21. September bis 16. Oktober ist eine interne Erprobungsplanung. Sie ist keine Aussage über gesetzliche Anwendungsfristen.

## 2 Separat angebotene Entwurfsfunktion

Version 2.0 kann anhand der vollständigen Textschicht einen Sachverhalt, eine Zuordnung zu Anspruchsmerkmalen und einen begründeten Entscheidungsvorschlag erzeugen. Die Funktion wählt automatisch Textstellen aus, die das Modell für den jeweiligen Streitpunkt relevant hält. Nicht übernommene Textstellen werden nicht gelöscht, erscheinen aber nur im gesondert zu öffnenden Originalfenster. Die für Anwaltskanzleien vertriebene Schwesterkonfiguration erstellt parteiliche Argumentationsentwürfe; diese Konfiguration ist nicht Gegenstand der gerichtlichen Projektinstallation.

Im Angebot sind zwölf Testlizenzen enthalten. Die Entwurfsfunktion ist mit dem gesonderten Schalter draft_enabled ausgestattet. Eine dauerhafte Abschaltung kann der zentrale Administrator technisch setzen. Eine Anleitung, den Schalter nicht zu betätigen, verändert seine technische Verfügbarkeit nicht. Das Modul darf derzeit keine Dateien an das gerichtliche Fachverfahren zurückschreiben.

## 3 Noch nicht beauftragte Erweiterung

Der Vorabstand 2.1 enthält eine Kennzeichnung widersprüchlicher Zeugenaussagen. Im Vertriebsgespräch wurde sie als Konsistenzhilfe gezeigt. Die angebotene Demonstration erzeugt zusätzlich einen numerischen Zuverlässigkeitswert pro Aussageperson. Ob diese Funktion in die spätere Ausschreibung aufgenommen wird, ist offen. Für die Testinstanz wurde am 1. Oktober nur eine zeitlich begrenzte Demonstrationsfreigabe vereinbart. Eine fachliche oder rechtliche Abnahme liegt nicht vor.

Friedemann Sturm
Leitung Projekte'''),
record('07_Protokoll_Projektbesprechung.docx','Besprechung zum Testlauf am 6 Oktober','Dr. Noura Seidlein | Projektstelle Justizdigitalisierung | Federweg 12 | 07745 Jena','''Teilnehmende: Dr. Noura Seidlein, Projektleitung; Kunigunde Häckel, Geschäftsstelle; Dr. Mira Schwanfelder, fachliche Begleitung; Emil Yalçın, IT.
Protokoll vom 6. Oktober 2026, 14:00 bis 15:10 Uhr

## 1 Tatsächlich genutzte Funktionen

Frau Häckel hat den Paginierungstest mit 30 synthetischen Akten abgeschlossen. Vier Register enthielten mindestens einen falschen Blattverweis. Die Originale blieben in allen 30 Akten zugänglich. Sie hat die fehlerhaften Verweise vor Ablage berichtigt und erklärt, dass ihr Zeitgewinn trotzdem spürbar sei. Die Zahl der Registerfehler sagt nicht, ob eine spätere materielle Entscheidung richtig wäre; solche Entscheidungen wurden in diesem Teiltest nicht erzeugt.

Dr. Schwanfelder hat für die Entwurfsfunktion dieselben 30 Akten verwendet. Bei sieben Akten fehlte im Vorschlag eine nach ihrer Durchsicht erhebliche Passage. Sie hatte zunächst nur bei 21 Akten sämtliche Originaldokumente im Testfenster geöffnet. Bei den übrigen neun prüfte sie gezielt die Fundstellen des Vorschlags. Eine zusätzliche Prüfung außerhalb der Anwendung ist nicht dokumentiert. Alle Akten sind erfunden; es wurde nichts in ein echtes Verfahren übernommen.

## 2 Widerspruch zur ersten Vertriebsantwort

Herr Sturm hatte am 29. September erklärt, das System treffe keine Auswahl. Die technische Darstellung vom 5. Oktober beschreibt dagegen einen Relevanzfilter. Frau Seidlein bittet um Aufklärung, ob damit nur die Textdarstellung oder auch der für den Entwurf berücksichtigte Akteninhalt begrenzt wird. Die bisherige Anbieterantwort betrifft diesen Unterschied nicht. Herr Yalçın kann bestätigen, dass der Nutzer das Original öffnen kann, nicht aber, dass er es vor jeder Übernahme wirklich liest.

## 3 Weitere Projektplanung

Bis zur Beratung der Projektgruppe bleibt der Zugang zu 2.1 auf zwei benannte Testkonten beschränkt. Herr Yalçın soll dokumentieren, ob die Demonstrationsfreigabe technisch ausläuft und ob die Zuverlässigkeitswerte exportiert wurden. Frau Häckel benötigt keine Bewertung von Personen für die Registerarbeit. Sie bittet darum, den bisherigen Paginierungsablauf in der rechtlichen Beurteilung nicht mit der neuen Funktion zu vermischen.

Die Projektgruppe hat noch keine Freigabe für eine Verwendung in echten Gerichtsakten erteilt. Ein möglicherweise späterer Einsatz ab Januar 2028 ist ein Planungswunsch, kein bereits beschlossener Termin. Dr. Seidlein benötigt zunächst einen ausformulierten Vermerk und ein Schreiben an den Anbieter.'''),
record('10_Stellungnahme_Arbeitsentwurf.docx','Arbeitsentwurf zur Einstufung von Aktenklar','Dr. Noura Seidlein | interner Entwurf | Stand 7. Oktober 2026','''## 1 Gegenstand und Tatsachenstand

Dieser Entwurf soll die Versionen 1.4, 2.0 und die Demonstration 2.1 getrennt beurteilen. Die Projektstelle beschafft das Werkzeug für eine spätere gerichtliche Verwendung. Die Testakten sind synthetisch; echte richterliche Entscheidungen oder eine Freigabe für Echtdaten sind nicht Bestandteil dieses Vorgangs. Der endgültig vorgesehene Nutzerkreis und die technische Abschaltung nicht beschaffter Module sind noch schriftlich zu bestätigen.

## 2 Vorläufiger Ansatz

Für die eigentliche Rechtsprüfung ist zu klären, welche Funktion das jeweilige System im Verhältnis zur menschlichen Arbeit übernimmt. Die Paginierung soll Akten erschließen. Die Entwurfsfunktion verarbeitet den Sachverhalt und schlägt eine rechtliche Bewertung vor. Die Demonstration versieht Aussagen zusätzlich mit personenbezogenen Zuverlässigkeitswerten. Das Wort Assistenz beschreibt diese Unterschiede nicht hinreichend.

Die Anbietererklärung vom 29. September verweist auf menschliche Letztentscheidung. Ob damit alle Voraussetzungen einer Ausnahme erfüllt sind, muss anhand des aktuellen Wortlauts und der tatsächlichen Auswahlwirkung geprüft werden. Der Relevanzfilter und die nur teilweise vollständige Originalsichtung sind dabei einzubeziehen. Eine fehlende echte Gerichtsentscheidung im Test besagt noch nicht, welche Zweckbestimmung für den späteren Betrieb vereinbart werden soll.

## 3 Auszuarbeitende rechtliche Bewertung

Die abschließende Fassung soll den einschlägigen Anhang-III-Untertatbestand, die vier alternativen Bedingungen des Artikels 6 Absatz 3 und eine etwaige persönliche Bewertung getrennt subsumieren. Der Produktpfad ist anhand der tatsächlichen Softwarefunktion zu prüfen, ohne ihn aus der bloßen Nutzung in der Justiz herzuleiten. Die Bewertung der Kanzleikonfiguration ist nur soweit nötig zur Abgrenzung heranzuziehen; sie ist nicht die gerichtliche Testinstanz.

Die zeitliche Anwendung der einzelnen Pflichten und ein etwaiges Bestandsrecht sind noch einzutragen. Der geplante Start im Januar 2028 ist nicht bestätigt. Es fehlen weiterhin die endgültige Zweckbeschreibung, die Erklärung zum Relevanzfilter und eine belastbare Beschreibung des Umfangs der menschlichen Prüfung. Die Rechtsberatung soll die offenen Punkte als konkrete Tatsachenfragen in ein Anbieteranschreiben übersetzen und die bereits entscheidbaren Teile vollständig begründen.

## 4 Geplante Entscheidungsvorlage

Die Projektleitung möchte für jede Version wissen, welcher Einsatz auf welcher Tatsachengrundlage beurteilt wurde und welche Änderung eine erneute Prüfung auslöst. Eine allgemeine Freigabe aller Module ist nicht beantragt. Die endgültige Empfehlung sowie die verantwortliche Zeichnung werden nach fachlicher und rechtlicher Beratung ergänzt. Bis dahin bleibt dies ein Arbeitsentwurf ohne Außenwirkung.'''),
record('03_Anleitung_Register.pdf','Aktenklar 1 4 Registeransicht','Aktenklar Software GmbH | technische Anleitung | 20. September 2026','''## 1 Eingabe und Ausgabe

Der Layoutdienst liest Scans, ordnet Dokumentanfänge zu und erstellt ein Register mit Blattverweisen. Die Funktion verwendet ein trainiertes Modell zur Erkennung von Layout und Textgrenzen. Sie bewertet weder die Schlüssigkeit eines Anspruchs noch die Zuverlässigkeit einer Person. Rechtsbegriffe werden als Überschriften übernommen, ohne daraus eine eigene rechtliche Folgerung zu erzeugen.

Das Originalfenster bleibt mit allen importierten Dateien zugänglich. Fehlerhafte Blattverweise sind vor Weitergabe des Registers zu berichtigen. Die Korrektur wird mit Konto und Zeitstempel gespeichert. Das Register ist eine Navigationshilfe; der zugrunde liegende Scan wird nicht überschrieben. Der Export enthält sämtliche Registerzeilen, auch wenn die Erkennung unsicher war.

## 2 Technische Grenzen

Handschriftliche Randbemerkungen und schlecht lesbare Faxköpfe können unzutreffende Dokumentanfänge auslösen. Der Status unsicher zeigt einen technischen Erkennungsbefund und keine rechtliche Wertung. Ein Link auf die falsche Seite kann korrigiert werden, ohne eine Quelle neu hochzuladen. Für den Testauftrag ist vorgesehen, dass die Geschäftsstelle die Blattzahlen anhand des Originals prüft.

Diese Anleitung gilt ausschließlich für Version 1.4. Die Schaltfläche Entscheidungsentwurf gehört zu einer gesonderten Funktion. Ihre Ausgaben, Relevanzfilter und personenbezogenen Kennzeichnungen werden durch diese Anleitung nicht beschrieben. Eine Bedienung der Registerfunktion bestätigt keine Prüfung der später hinzugefügten Module.'''),
record('05_Auszug_Entwurfstest.pdf','Auszug aus dem Test der Entwurfsfunktion','Dr. Mira Schwanfelder | fachliche Testbegleitung | 5. Oktober 2026','''## 1 Testakte J 17

Die synthetische Akte betrifft die Lieferung einer Werkbank. In Dokument 8 erklärt die Käuferin, dass sie eine Ersatzlieferung angenommen habe; Dokument 11 enthält die abweichende Lieferscheinnummer. Der erzeugte Entwurf würdigt die erste Mängelanzeige ausführlich, erwähnt aber die Ersatzlieferung nicht. Im Auswahlfenster wurden fünf von insgesamt neun importierten Dokumenten als relevant angezeigt. Die beiden genannten Dokumente waren nicht unter diesen fünf.

Die Schaltfläche Alle Originale öffnet sämtliche Dateien. Der Testbericht bestätigt deshalb keine technische Löschung. Ungeklärt ist, ob das Modell die ausgeblendeten Dokumente vor Erzeugung des Entwurfs vollständig berücksichtigte oder ob bereits seine Eingabe auf die Auswahl begrenzt war. Der Anbieter hat dazu einen Ablaufnachweis angekündigt. Die bisherige Aussage, alle Daten seien im System vorhanden, beantwortet diese Frage nicht.

## 2 Vergleich der Funktionsstände

Die frühere Paginierung zeigt in derselben Testakte alle neun Dokumente. Version 2.0 erzeugt zusätzlich einen rechtlichen Entscheidungsvorschlag. In der Demonstration 2.1 erscheint bei der erfundenen Zeugin Elfriede Schnackenberg der Wert 0,34 mit der Bezeichnung Zuverlässigkeit. Die angezeigte Begründung verweist auf angeblich wechselnde Formulierungen in ihren Aussagen. Eine fachliche Validierung dieser Kennzeichnung wurde nicht durchgeführt.

Die Testbegleitung hat keine dieser Ausgaben in ein gerichtliches Verfahren übertragen. Ihr Befund betrifft die beobachtete Funktion und den Bedarf an Aufklärung. Er ist keine rechtliche Zertifizierung und keine Aussage über die Glaubwürdigkeit einer wirklichen Person.'''),
record('11_Angebot_Projektaufwand.pdf','Angebot für den nächsten Testabschnitt','Aktenklar Software GmbH | Friedemann Sturm | 6. Oktober 2026','''An die Projektstelle Justizdigitalisierung Jena, Federweg 12, 07745 Jena
Angebot AK-26-129, gültig bis 23. Oktober 2026

## 1 Leistungspositionen

Für zwölf Testkonten werden monatlich 180,00 EUR netto je Konto berechnet. Die Einrichtung des zentralen Modulschalters und eines Exportprotokolls wird einmalig mit 1.600,00 EUR netto angeboten. Eine eintägige Schulung für die Projektgruppe kostet 950,00 EUR netto. Die Preise enthalten keine rechtliche Begutachtung und keine fachliche Validierung richterlicher Entscheidungsentwürfe.

Für drei Monate ergeben sich 6.480,00 EUR netto Lizenzentgelt. Einrichtung und Schulung zusammen betragen 2.550,00 EUR netto. Der Nettogesamtbetrag lautet 9.030,00 EUR; bei 19 Prozent Umsatzsteuer beträgt die Steuer 1.715,70 EUR und der Bruttobetrag 10.745,70 EUR. Dies ist ein Angebot und keine fällige Rechnung. Ein Auftrag ist bislang nicht eingegangen.

## 2 Umfang und Vorbehalte

Die vollständige Abschaltung von 2.1 ist Bestandteil der Einrichtung, wenn die Projektstelle sie beauftragt. Die Freigabe oder Ablehnung einzelner Nutzungszwecke verbleibt beim Auftraggeber nach seiner Prüfung. Der angebotene Export zeigt Konten, Zeitpunkte und geöffnete Dokumente; er erfasst keine außerhalb der Anwendung gelesenen Papierakten. Der nächste Testabschnitt soll ausschließlich mit den vereinbarten synthetischen Akten durchgeführt werden.

Friedemann Sturm
Leitung Projekte''')],
emails=[
email('01_Auftrag.eml','Aktenklar bitte die drei Fassungen getrennt beurteilen','Dr. Noura Seidlein <noura.seidlein@justizprojekt.example>','Dr. Salome Fritsch <salome.fritsch@kanzlei-fritsch.example>','2026-10-09T08:15:00+02:00','''Sehr geehrte Frau Dr. Fritsch,

bitte erstellen Sie für unsere Projektgruppe einen ausformulierten Einstufungsvermerk zur KI-Verordnung. Wir benötigen eine getrennte Betrachtung der Registerfunktion 1.4, der Entwurfsfunktion 2.0 und der Demonstration 2.1. Im Anhang finden Sie den technischen Testauszug. Unser Arbeitsentwurf liegt ebenfalls im Ordner; bitte bearbeiten Sie ihn weiter und kennzeichnen Sie konkret, welche Tatsachen noch fehlen.

Außerdem benötigen wir einen Brief an Aktenklar, der die offenen Fragen so stellt, dass die Entwickler sie beantworten können. Ich möchte keine allgemeine Liste aller denkbaren Compliance-Unterlagen. Im Mittelpunkt stehen die Auswahl der Dokumente, die persönlichen Zuverlässigkeitswerte und der behauptete lediglich vorbereitende Charakter. Bitte prüfen Sie ausdrücklich die Ausnahme nach Artikel 6 Absatz 3, ohne sie vorwegzunehmen.

Ein Einsatz mit echten Gerichtsakten wurde noch nicht genehmigt. Die Projektgruppe diskutiert Januar 2028, hat sich aber nicht entschieden. Das Angebot soll erst nach unserer Besprechung nächste Woche bestellt werden. Bitte senden Sie weder dem Anbieter noch einer Behörde etwas in unserem Namen. Wir brauchen zunächst die intern prüfbaren Entwürfe.

Mit freundlichen Grüßen
Dr. Noura Seidlein''',['05_Auszug_Entwurfstest.pdf']),
email('04_Vertrieb_Ausnahme.eml','Unsere Assistenz bleibt aus unserer Sicht vorbereitend','Friedemann Sturm <friedemann.sturm@aktenklar.example>','Dr. Noura Seidlein <noura.seidlein@justizprojekt.example>','2026-09-29T16:40:00+02:00','''Sehr geehrte Frau Dr. Seidlein,

unser Vertrieb geht davon aus, dass die Funktionen insgesamt vorbereitend bleiben. Das System spricht kein Urteil; jeder Text wird erst durch die zuständige Person übernommen. Deshalb hatten wir im Gespräch eine Ausnahme für das Gesamtprodukt angesprochen. Eine gesonderte schriftliche rechtliche Bewertung für die drei Versionen können wir Ihnen derzeit noch nicht übersenden.

Wir haben ein internes Qualitätsmanagement und planen eine Zertifizierung. Ein Systemzertifikat für den gerichtlichen Entwurfsprozess liegt nicht vor. Eine Registrierungsnummer können wir ebenfalls nicht nennen; wir gingen bisher davon aus, dass eine Ausnahme eine solche Registrierung entbehrlich mache. Ich werde diese Aussage mit unserer Rechtsabteilung klären.

Zu Ihrer Frage nach der Auswahl: Sämtliche Akten bleiben gespeichert. Die Liste im Vorschaufenster dient aus Sicht des Produktteams der besseren Übersicht. Ob der Entwurfsdienst zuvor alle Texte in gleicher Weise verarbeitet, muss die Entwicklung beantworten. Ich möchte dazu keine falsche technische Zusicherung abgeben. Bitte behandeln Sie diese Nachricht bis zu dieser Klärung als Vertriebsstand, nicht als verbindliche Abnahmeerklärung.

Mit freundlichen Grüßen
Friedemann Sturm'''),
email('08_IT_Zugriffsprotokoll.eml','Export zählt Öffnungen innerhalb der Anwendung','Emil Yalçın <emil.yalcin@justizprojekt.example>','Dr. Noura Seidlein <noura.seidlein@justizprojekt.example>','2026-10-07T11:05:00+02:00','''Guten Tag Frau Dr. Seidlein,

anbei die Arbeitsmappe mit den beiden Teiltests und den aus dem Angebot übernommenen Kosten. Bitte nicht aus der Zahl der vollständigen Öffnungen schließen, dass die übrigen Akten außerhalb der Anwendung niemand gelesen habe. Genau das misst der Export nicht. Frau Schwanfelder sagte mir allerdings, dass sie für neun Entwürfe zunächst nur die vorgeschlagenen Fundstellen geöffnet hatte.

Die Demonstrationsfunktion 2.1 war vom 1. bis 6. Oktober nur für zwei Testkonten sichtbar. Das allgemeine Teamkonto konnte sie nicht aufrufen. Die Sperre seit gestern ist technisch aktiv; ich habe den Schalter und einen fehlgeschlagenen Aufruf dokumentiert. Der Anbieter muss noch bestätigen, ob die entstandenen Zuverlässigkeitswerte im Hintergrund weiter für die nächste Sitzung verwendet werden.

Für 1.4 ist kein automatischer Versand eingerichtet. Auch 2.0 kann nicht in das Fachverfahren schreiben. Ein kopierter Entwurf kann aber außerhalb der Anwendung weiterbearbeitet werden. Unser Not-Aus verhindert neue Modellaufrufe, entfernt keine bereits exportierten Dateien. Bitte berücksichtigen Sie diese Grenze bei jeder Formulierung zur menschlichen Kontrolle.

Freundliche Grüße
Emil Yalçın''',['09_Testlauf_und_Aufwand.xlsx']),
email('12_Fachliche_Rueckfrage.eml','Paginierung bitte weiter nutzbar prüfen','Kunigunde Häckel <kunigunde.haeckel@justizprojekt.example>','Dr. Noura Seidlein <noura.seidlein@justizprojekt.example>','2026-10-08T15:25:00+02:00','''Sehr geehrte Frau Dr. Seidlein,

für die Geschäftsstelle ist die Registerarbeit der eigentliche Nutzen. Ich habe alle vier falschen Blattverweise korrigiert. Im Register steht keine persönliche Bewertung, und keine Akte wurde wegen eines Modellwerts ausgeblendet. Wenn die neue Entwurfsfunktion rechtlich schwieriger ist, sollten wir bitte prüfen lassen, ob 1.4 davon tatsächlich abtrennbar bleibt.

Mir ist wichtig, dass im späteren Schreiben nicht steht, wir hätten sämtliche Entwürfe umfassend gelesen. Das war nicht meine Aufgabe und kann ich nicht bestätigen. Ich habe nur die Registeransicht geprüft. Frau Schwanfelder hat die materiellen Entwürfe angesehen. Über eine Verwendung in echten Verfahren kann ich ohnehin nicht entscheiden.

Die Namen im Testauszug gehören zu den erfundenen Akten. Eine Schulung würde ich gern mitmachen, sie ist aber noch nicht gebucht. Bitte fragen Sie den Anbieter auch, ob die Korrekturen in der nächsten Version übernommen werden oder ob wir die Blattverweise nach jedem Update vollständig neu kontrollieren müssen.

Mit freundlichen Grüßen
Kunigunde Häckel''')],
chatfile='06_Projektchat.txt',chat='''Projektchat Aktenklar, exportiert am 7. Oktober 2026. Ausschließlich Projektbeteiligte. Bezug JA-2026-104.

05.10.2026 09:05, Mira: Im Fall J17 fehlt mir die Ersatzlieferung. Das Original ist da, im Vorschlag taucht es nicht auf.
05.10.2026 09:08, Emil: Alle Originale öffnet neun Dateien. In der Vorschau sehe ich fünf. Ich weiß noch nicht, was der Modellaufruf bekommt.
05.10.2026 09:11, Noura: Bitte genau diesen Unterschied dokumentieren. Nicht vorschnell schreiben, die Dateien seien gelöscht.
05.10.2026 09:15, Kunigunde: In 1.4 steht die Ersatzlieferung im Register auf Blatt 47. Der Link war korrekt.
05.10.2026 09:21, Mira: Der neue Wert Zuverlässigkeit 0,34 ist eine andere Sache. Ich kann nicht erkennen, ob nur Textwidersprüche oder die Person bewertet werden.
05.10.2026 09:23, Noura: Bitte Screenshot und den vollständigen Text sichern. Keine echten Namen eingeben.
06.10.2026 16:02, Emil: Schalter für 2.1 jetzt gesperrt. Ein bereits exportiertes Ergebnis verschwindet dadurch natürlich nicht.
06.10.2026 16:04, Noura: Danke. Wir klären die Fassungen getrennt. Noch keine Bestellung der zusätzlichen Monate.
07.10.2026 10:35, Mira: Ich habe 21 Akten vollständig geöffnet. Bei neun zunächst nur die vorgeschlagenen Belegstellen. Bitte so in die Tabelle.
07.10.2026 10:39, Emil: Erfasst. Die Tabelle enthält keine Wertung, ob das ausreicht.
''',xlsx='09_Testlauf_und_Aufwand.xlsx',
rows=[['Register 1.4',30,4,30],['Entwurf 2.0',30,7,21]],costs=[['Testkonten je Monat',12,180,3],['Einrichtung',1,1600,1],['Schulung',1,950,1]],
notes=['B6:B7: synthetische Akten je Teiltest; dieselben 30 Akten, nicht 60 verschiedene Personen.','C6: falsche Blattverweise; C7: mindestens eine erhebliche Passage im Entwurf nicht genannt.','D6:D7: vollständig im Testfenster geöffnete Originalakten; keine Messung außerhalb des Systems.','Quelle: 07_Protokoll_Projektbesprechung.docx und 08_IT_Zugriffsprotokoll.eml.','Kostenquelle: 11_Angebot_Projektaufwand.pdf. Noch nicht bestellt, keine fällige Rechnung.']),

dict(slug='ki-hochrisiko-medizinprodukt-saalfeld',title='Saalfeld Bildspur zwischen Archiv und Befundhilfe',reference='BS-2026-219',
intro='Das fiktive Klinikum Morgenhöhe in Saalfeld prüft eine neue Bildassistenz. Die Anbieterbeschreibung spricht von Komfort; ein Labortest zeigt einen übersehenen Bildbefund. Die Akte enthält getrennte Versionen, eine offene medizinproduktrechtliche Einordnung, einen Testdatensatz und eine noch nicht beschlossene Notfalloption.',
docs=[
record('02_Produktbeschreibung.docx','Bildspur Funktionsstände und Zweck','Dr. Livia Knörlein | Bildspur Medizintechnik GmbH | Messweg 6 | 07318 Saalfeld','''An das Klinikum Morgenhöhe, Dr. Hanno Wölfel, IT und Medizintechnik
Stand 22. September 2026, Projekt BS-2026-219

## 1 Archivfunktion 1 0

Bildspur 1.0 erkennt Dokumenttypen und führt vorhandene Untersuchungsbezeichnungen aus dem Archiv in einer Suchliste zusammen. Die trainierte Texterkennung wertet die Bildinhalte nicht diagnostisch aus. Der Dienst trifft keine Aussage über einen Befund und beeinflusst nicht die Reihenfolge der ärztlichen Befundung. Für die Darstellung diagnostischer Bilder verwendet die Klinik ihr gesondertes vorhandenes Befundsystem.

## 2 Bildassistenz 2 0

Die neue Version verarbeitet CT-Bildserien nach ihrer Rekonstruktion. Ein trainiertes Bildmodell kennzeichnet Bildabschnitte, in denen nach seinem Muster eine auffällige Veränderung vorliegen könnte. Die Standardansicht zeigt diese Abschnitte zuerst. Unmarkierte Bilder bleiben über die Schaltfläche Gesamtserie erreichbar. Die Funktion ist für die Unterstützung bei Verlaufskontrollen in der ambulanten radiologischen Routine vorgesehen. Ein ärztlicher Befund wird vom System nicht unterschrieben oder selbst versandt.

Das Produktteam bezeichnet die neue Darstellung intern als Komfortmodus. Die medizinische Zweckbeschreibung lautet gegenwärtig Unterstützung der ärztlichen Bewertung auffälliger Bildveränderungen. Diese Formulierungen werden von der Rechtsabteilung noch abgeglichen. Für 2.0 liegt keine abschließend unterzeichnete EU-Konformitätserklärung in diesem Projektordner. Der bisher übersandte Nachweis betrifft unser organisationsbezogenes Qualitätsmanagement, nicht die Freigabe dieser konkreten Funktion.

## 3 Noch offene Optionen

Das Vertriebsteam hat zusätzlich eine mögliche Notfallvariante ab 2027 angesprochen. Sie würde Untersuchungen anhand eines vermuteten akuten Befunds priorisieren. Diese Option ist weder Bestandteil der Testlizenz noch technisch aktiviert. Der jetzige Test findet mit 40 synthetischen Bildserien in einem vom Klinikbetrieb abgetrennten Laborbereich statt. Ein Zugriff auf aktuelle Patientenakten ist nicht eingerichtet.

Dr. Livia Knörlein
Produktverantwortliche'''),
record('07_Klinische_Besprechung.docx','Besprechung zu Bildauswahl und Rückfallweg','Dr. Ottilie Wagenblast | Radiologie Klinikum Morgenhöhe | Hangblick 18 | 07318 Saalfeld','''Besprechung am 7. Oktober 2026, 13:30 Uhr
Teilnehmende: Dr. Ottilie Wagenblast, Radiologie; Dr. Hanno Wölfel, IT; Berfin Degenhardt, Qualitätsmanagement; Dr. Livia Knörlein, Hersteller.

## 1 Beobachtung im Labor

Vierzig synthetische Serien wurden mit 2.0 geprüft. Vier Serien enthielten nach dem vorab festgelegten Referenzbefund eine auffällige Veränderung. Das Modell markierte drei davon und übersah im Fall S17 die entsprechende Bildfolge. Drei zusätzliche unauffällige Serien erhielten einen auffälligen Marker. Die Testgruppe leitete keine Behandlung daraus ab. Es wurden keine echten Patienten geschädigt oder untersucht.

Bei S17 war die vollständige Serie über eine zweite Schaltfläche erreichbar. Dr. Wagenblast konnte den Referenzbefund nach Öffnung der Gesamtserie nachvollziehen. Im ersten Testlauf hatten die Nutzer aber nur bei 28 von 40 Serien diese Gesamtansicht geöffnet. Eine verpflichtende technische Bestätigung der vollständigen Sichtung existiert noch nicht. Die Anleitung verlangt eine eigenständige ärztliche Befundung; ob der geplante Arbeitsablauf sie tatsächlich sicherstellt, bleibt zu prüfen.

## 2 Bedeutung der unterschiedlichen Angaben

Die Herstellerin erklärt, dass die Funktion keine Therapie anordne. Dr. Wagenblast weist darauf hin, dass eine dauerhaft ausgelassene Auffälligkeit dennoch die ärztliche Beurteilung beeinflussen könne. Wie schwer eine solche Folge im vorgesehenen Routineeinsatz werden könnte, ist medizinisch noch nicht bewertet. Die klinische Leiterin möchte vor einer Einführungsentscheidung eine dokumentierte Gefährdungsanalyse und eine eindeutige Anleitung erhalten.

## 3 Noch keine Einführung

Dr. Wölfel bestätigt, dass das Testnetz keinen Zugriff auf den Produktivviewer hat. Der Hersteller soll den Vorschlag für eine neue Standardsicht liefern, in der zunächst alle Bilder und erst anschließend die Markierungen erscheinen. Die Klinik hat diesen Vorschlag noch nicht abgenommen. Der Beschaffungswunsch für Januar 2027 ist vorläufig; eine produktive Freigabe wurde nicht erteilt.

Die geplante Notfallpriorisierung wurde nur als spätere Option angesprochen. Die Teilnehmer sind sich einig, dass sie nicht stillschweigend in die jetzige Prüfung einbezogen werden darf. Die dafür benötigten rechtlichen und technischen Nachweise wären gesondert zu behandeln.'''),
record('10_Stellungnahme_Arbeitsentwurf.docx','Arbeitsentwurf zum Produktpfad von Bildspur','Berfin Degenhardt | Klinikum Morgenhöhe | 8. Oktober 2026','''## 1 Auftrag und Gegenstand

Dieser Entwurf bereitet eine interne Entscheidung über Bildspur 2.0 vor. Die Archivfunktion 1.0 und die noch nicht beauftragte Notfalloption müssen davon getrennt bleiben. Das Klinikum benötigt eine begründete Prüfung, ob der vorgesehene Einsatz einen Hochrisikopfad der KI-Verordnung eröffnet. Ein CE-Zeichen oder eine rechtliche Freigabe wird mit diesem Entwurf nicht erzeugt.

## 2 Bisherige Tatsachen

Die Herstellerin beschreibt die neue Funktion als Unterstützung der ärztlichen Bewertung von CT-Bildern, verwendet aber zugleich den Ausdruck Komfortmodus. Der Laborversuch zeigt eine nicht markierte Auffälligkeit bei S17. Die Gesamtserie blieb zugänglich, wurde jedoch nicht in allen Testvorgängen geöffnet. Welche Gesundheitsfolgen ein entsprechendes Übersehen im späteren Routineeinsatz haben könnte, ist bislang nicht schriftlich beurteilt.

Die Produktklasse und der einschlägige Weg der vorgeschriebenen Konformitätsbewertung sind offen. Der Hersteller hält eine bestimmte Klasse für wahrscheinlich, hat seine vollständige Begründung aber noch nicht vorgelegt. Das organisationsbezogene Qualitätsmanagementdokument benennt keine Freigabe der Version 2.0. Der Beschaffungsordner enthält daher noch keinen vollständigen Produktnachweis.

## 3 Auszuarbeitende Prüfung

Die Rechtsberatung soll beide Voraussetzungen des Artikels 6 Absatz 1 und die aktuellen Absätze 1a bis 1c anhand der Belege getrennt prüfen. Insbesondere ist zu erläutern, wie die behauptete nicht sicherheitsrelevante Unterstützung und die mögliche Gefährdung bei Ausfall oder Fehlfunktion zueinander stehen. Eine allgemeine vorbereitende Funktion darf nicht ohne Prüfung als Ausnahme für jedes Produkt behandelt werden.

Zusätzlich soll geklärt werden, ob der tatsächliche Routineeinsatz einen konkreten Anhang-III-Tatbestand erfüllt. Die nur diskutierte Notfalloption darf nicht als aktivierter Funktionsbestand zugrunde gelegt werden. Für den späteren Einsatz ist ein normbezogener Zeitplan erforderlich. Eine noch ausstehende Anwendung einzelner KI-Pflichten ersetzt keine Untersuchung bereits geltender medizinproduktrechtlicher oder sonstiger Anforderungen.

## 4 Nächste Dokumente

Der Einstufungsvermerk soll die offenen technischen und medizinischen Fragen präzise benennen. Ein getrenntes Schreiben an die Herstellerin soll Zweckbestimmung, Klassenbegründung, vorgeschriebenen Bewertungsweg, Version und Nachweisumfang anfordern. Die Empfehlung zum Starttermin wird erst nach den entscheidenden Antworten ergänzt. Der derzeitige Entwurf ist intern, unvollständig und nicht zur Übersendung als behördliche Erklärung vorgesehen.'''),
record('03_Technischer_Fehlertest.pdf','Bildspur Laborbeobachtung S17','Klinikum Morgenhöhe | Testlabor | Bericht vom 5. Oktober 2026','''## 1 Prüfumgebung

Der Test verwendet 40 synthetische CT-Serien mit einem vorab festgelegten Referenzbefund. Die Serien sind weder Bilder wirklicher Patienten noch eine repräsentative klinische Validierungsstichprobe. Der Versuch prüft die Darstellung und die Reaktion auf ausgewählte Bildmuster. Seine Kennzahlen dürfen nicht als Nachweis einer allgemeinen diagnostischen Leistungsfähigkeit verwendet werden.

## 2 Beobachtung S17

Die Referenzbeschreibung enthält eine auffällige Veränderung in einer hinteren Bildfolge. Version 2.0 erzeugte dort keinen Marker. In der Standardansicht waren zuerst andere markierte Bildbereiche sichtbar. Nach Auswahl der Gesamtserie konnte Dr. Wagenblast die betreffende Folge öffnen. Der technische Fehler lag daher nicht in einer vollständigen Löschung des Bildmaterials. Welche Eingabeschritte oder Modellmerkmale zum fehlenden Marker führten, wurde noch nicht untersucht.

Im geplanten Arbeitsablauf soll jede Ärztin und jeder Arzt die vollständige Serie eigenständig befunden. Im Test ist das nur für 28 von 40 Vorgängen im Viewer protokolliert. Die übrigen zwölf Einträge belegen keine vollständige Öffnung im Testfenster. Ob andere Wege genutzt wurden, ist nicht erfasst. Für eine Einführung wird ein überprüfbarer Ablauf benötigt; die bloße Zugänglichkeit beweist keine tatsächliche Sichtung.

## 3 Weiteres Vorgehen

Die Entwickler erhalten den synthetischen Datensatz und die Ereignisprotokolle im vereinbarten Testbereich. Der Zugriff auf echte Patientendaten wird dadurch nicht eröffnet. Ein neuer Test muss dieselben Ausfallbedingungen und den vorgeschlagenen Rückfallweg erfassen. Bis zu dessen Auswertung bleibt der Befund S17 als offener Fehler im Projekt bestehen.'''),
record('05_Herstellerstand_Produktrecht.pdf','Stand der Unterlagen für Bildspur 2 0','Dr. Livia Knörlein | Bildspur Medizintechnik GmbH | 6. Oktober 2026','''An das Klinikum Morgenhöhe
Bezug BS-2026-219

## 1 Medizinische Zweckbestimmung

Unser derzeitiger Entwurf der Zweckbeschreibung nennt die Unterstützung ärztlicher diagnostischer Entscheidungen bei radiologischen Verlaufskontrollen. Die interne Bezeichnung Komfortmodus ist nicht als eigenständige regulatorische Zweckbestimmung gedacht. Die Rechtsabteilung prüft gegenwärtig, welche medizinproduktrechtliche Klasse sich aus den vorgesehenen Entscheidungen und ihren möglichen Folgen ergibt. Im Projektgespräch wurde Klasse IIa als Arbeitshypothese genannt; eine abschließende Begründung ist damit nicht verbunden.

## 2 Bewertungsweg

Wir prüfen insbesondere die Software-Regel 11 und den hieraus folgenden Bewertungsweg. Die bislang an Sie gesandte Darstellung unseres Qualitätsmanagements betrifft die Organisation. Sie ersetzt keine Bescheinigung einer benannten Stelle für die neue Bildassistenz. Eine verbindliche Erklärung, dass gerade Version 2.0 von einem abgeschlossenen Verfahren erfasst sei, können wir zum heutigen Datum nicht abgeben. Die genaue Verfahrenszuordnung und ein möglicher Antrag werden noch bearbeitet.

## 3 Angebotene Zusammenarbeit

Wir werden die vollständige Klassenbegründung, den vorgesehenen Verfahrensweg und die aktuelle Gebrauchsanleitung nachreichen. Die medizinische Beurteilung möglicher Folgen eines übersehenen Bildbefunds bedarf der fachlichen Ergänzung. Die Unterschrift einer Ärztin oder eines Arztes unter einem späteren Befund ist aus unserer Sicht ein wichtiger Kontrollschritt, beantwortet aber nicht alle technischen Fragen des Ausfallpfads.

Die Notfallpriorisierung ist eine separate Entwicklung. Der Laborvertrag mit Ihrer Klinik enthält diese Funktion nicht. Bis zu einer schriftlichen Erweiterung dürfen die Testzugänge nicht für die Reihenfolge akuter Notfalluntersuchungen eingesetzt werden.

Dr. Livia Knörlein'''),
record('11_Kostenangebot.pdf','Angebot für die Fortführung des Bildspur Tests','Bildspur Medizintechnik GmbH | Angebot BS-26-311 | 7. Oktober 2026','''An das Klinikum Morgenhöhe, Hangblick 18, 07318 Saalfeld

## 1 Leistungsumfang und Preis

Für sechs Testzugänge über vier Monate beträgt das Entgelt je Zugang und Monat 240,00 EUR netto. Daraus ergeben sich 5.760,00 EUR netto. Für die Einbindung in das abgetrennte Testnetz werden einmalig 2.400,00 EUR netto angeboten. Zwei Schulungstermine kosten jeweils 750,00 EUR netto, zusammen 1.500,00 EUR netto. Die Nettosumme beträgt 9.660,00 EUR. Bei 19 Prozent Umsatzsteuer kommen 1.835,40 EUR hinzu; der Bruttobetrag lautet 11.495,40 EUR.

Dies ist ein unverbindlicher Kalkulationsstand für die Entscheidung des Klinikums und keine Rechnung. Eine Bestellung ist noch nicht erfolgt. Die Durchführung einer vollständigen klinischen Validierung und die gegebenenfalls erforderliche Mitwirkung einer benannten Stelle sind im Preis nicht enthalten. Dafür werden Umfang und Aufwand erst nach Festlegung der endgültigen Zweckbestimmung vereinbart.

## 2 Technische Abgrenzung

Die Testzugänge erlauben keinen Zugriff auf den Produktivviewer und keinen Versand medizinischer Befunde. Der Hersteller unterstützt die Untersuchung der beobachteten Fehler mit synthetischen Serien. Eine Änderung der Darstellung, die sämtliche Originalbilder vor den Markierungen zeigt, soll gesondert angeboten werden. Die Klinik muss darüber entscheiden, bevor ein entsprechender Entwicklungsauftrag ausgelöst wird.

Die Notfalloption ist nicht enthalten. Aus der Laufzeit dieses Angebots folgt keine Aussage über gesetzliche Übergangsfristen oder eine Erlaubnis zum Patienteneinsatz. Ansprechpartnerin für eine Rückfrage ist Dr. Livia Knörlein unter livia.knoerlein@bildspur.example.''')],
emails=[
email('01_Auftrag.eml','Bitte Bildspur Produktpfad und geplanten Einsatz getrennt prüfen','Dr. Hanno Wölfel <hanno.woelfel@morgenhoehe.example>','Dr. Salome Fritsch <salome.fritsch@kanzlei-fritsch.example>','2026-10-09T09:20:00+02:00','''Sehr geehrte Frau Dr. Fritsch,

bitte erstellen Sie einen ausformulierten Vermerk zur Hochrisikoeinstufung unserer geplanten Bildassistenz. Der Hersteller spricht von Komfort, nennt in seiner Produktbeschreibung aber auch diagnostische Unterstützung. Den Laborbericht S17 finden Sie anbei. Unser internes Word-Dokument ist nur ein Arbeitsentwurf und soll durch Ihre Prüfung weitergeführt werden.

Wir brauchen insbesondere eine saubere Trennung zwischen der alten Archivfunktion, der neuen Bildassistenz und der nur besprochenen Notfalloption. Bitte prüfen Sie die Produktvoraussetzungen einschließlich der Frage, was der Hinweis auf menschliche Befundung tatsächlich trägt. Sagen Sie uns auch, welche konkreten Unterlagen wir noch anfordern müssen; eine umfassende technische Zertifizierung ist nicht Ihr Auftrag.

Wir planen einen Start im Januar 2027, haben ihn jedoch noch nicht freigegeben. Es gibt derzeit keine echten Patientendaten in der Testumgebung. Bitte erstellen Sie ein getrenntes Anbieteranschreiben und eine kurze interne Entscheidungsvorlage. Nichts soll bereits versandt, gemeldet oder bestellt werden. Über den weiteren Umfang entscheiden wir nach der Projektbesprechung.

Mit freundlichen Grüßen
Dr. Hanno Wölfel''',['03_Technischer_Fehlertest.pdf']),
email('04_Vertriebsargument.eml','Komfortfunktion und persönliche Befundung','Ottmar Seidel <ottmar.seidel@bildspur.example>','Dr. Hanno Wölfel <hanno.woelfel@morgenhoehe.example>','2026-10-02T14:35:00+02:00','''Sehr geehrter Herr Dr. Wölfel,

aus Vertriebssicht soll die Software vor allem Zeit beim Öffnen von Bildserien sparen. Sie unterschreibt nichts und ordnet keine Behandlung an. Ich hatte deshalb vorgeschlagen, den Ablauf als bloße Vorbereitung zu betrachten. Ob die neueren Sicherheitsvorschriften der KI-Verordnung hierzu eine genauere Einordnung verlangen, habe ich noch nicht mit unserer Rechtsabteilung abgestimmt.

Unser Qualitätsmanagement ist ausgebaut, und wir arbeiten mit externen Prüfern zusammen. Daraus wollte ich keine Aussage ableiten, dass bereits jede neue Version ein eigenes Systemzertifikat erhalten hätte. Frau Dr. Knörlein wird Ihnen den tatsächlichen Stand der Produktunterlagen gesondert mitteilen. Bitte legen Sie meiner Nachricht keine Nummer einer benannten Stelle oder eine bereits bestehende Produktzulassung zugrunde.

Die auf der Messe gezeigte Notfallpriorisierung war eine Entwicklungsansicht. In Ihren Testkonten ist sie nicht vorhanden. Eine Nutzung in der Notaufnahme würde ein eigenes Projekt erfordern. Für Ihre Entscheidung stelle ich gern einen Termin mit der Produktverantwortlichen und der klinischen Ansprechpartnerin ein.

Mit freundlichen Grüßen
Ottmar Seidel'''),
email('08_Testdaten.eml','Arbeitsmappe zum synthetischen Test und zu den Angebotskosten','Berfin Degenhardt <berfin.degenhardt@morgenhoehe.example>','Dr. Hanno Wölfel <hanno.woelfel@morgenhoehe.example>','2026-10-08T10:10:00+02:00','''Guten Tag Herr Dr. Wölfel,

anbei die Excel-Datei. Die Kennzahlen beziehen sich nur auf den kleinen synthetischen Laborversuch. Vierzig Serien mit vier Referenzauffälligkeiten erlauben keine Aussage über die klinische Sensitivität im Patienteneinsatz. Ich habe deshalb Zählwerte und Quotienten benannt und auf eine grüne Sicherheitsampel verzichtet. Bitte lassen Sie auch die juristische Einordnung nicht aus einer Prozentzahl ableiten.

Der Fehler S17 ist weiterhin offen. Ein zusätzlicher Test soll mit geänderter Startansicht durchgeführt werden; die Daten in der Arbeitsmappe zeigen noch die bisherige Version. Die Kostenseite entspricht dem Angebot vom 7. Oktober und darf nicht als bereits fällige Verbindlichkeit in die Buchhaltung übernommen werden.

Die klinische Leiterin konnte noch nicht schriftlich bestätigen, welche Folgen ein übersehener Befund im gesamten vorgesehenen Einsatzspektrum haben könnte. Diese Frage betrifft aus meiner Sicht die Produktbeschreibung und nicht nur die Schulung. Bitte nehmen Sie sie in das Anbieteranschreiben beziehungsweise die fachliche Nachforderung auf.

Freundliche Grüße
Berfin Degenhardt''',['09_Testdaten_und_Kosten.xlsx']),
email('12_Hersteller_Nachtrag.eml','Neue Gesamtansicht ist erst ein Vorschlag','Dr. Livia Knörlein <livia.knoerlein@bildspur.example>','Dr. Hanno Wölfel <hanno.woelfel@morgenhoehe.example>','2026-10-08T17:45:00+02:00','''Sehr geehrter Herr Dr. Wölfel,

wir können die vollständige Bildserie vor den Markierungen anzeigen lassen. Die Entwicklung hat dafür einen Entwurf vorgelegt, aber noch keinen freigegebenen Versionsstand ausgeliefert. Bitte behandeln Sie diesen Vorschlag nicht als bereits wirksame Sicherung in Version 2.0. Der neue Test soll auch prüfen, ob die Bedienung unter Zeitdruck tatsächlich verändert wird.

Die Klassenbegründung ist weiterhin in Bearbeitung. Ich bitte darum, weder unsere bisherige Arbeitshypothese noch das Organisationszertifikat als endgültigen Produktnachweis zu verwenden. Für den rechtlichen Vermerk können Sie festhalten, dass wir die Unterstützung diagnostischer Entscheidungen als Zweck prüfen und die möglichen Folgen einer Fehlentscheidung mit Fachleuten bewerten lassen.

Eine Einbindung in echte Befundungen ist durch den aktuellen Testvertrag nicht gestattet. Die Laborumgebung bleibt getrennt. Zum weiteren Termin melde ich mich nach der Abstimmung mit unserem Qualitätsmanagement. Ich kann heute noch keine Aussage machen, wann ein vollständiges Konformitätspaket für Ihre konkrete Konfiguration vorliegen wird.

Mit freundlichen Grüßen
Dr. Livia Knörlein''')],
chatfile='06_Testlabor_Chat.txt',chat='''Testlabor Bildspur, Export vom 8. Oktober 2026. Bezug BS-2026-219.

05.10.2026 08:40, Ottilie: S17 hat keinen Marker. In der vollständigen Serie ist die Referenzauffälligkeit sichtbar.
05.10.2026 08:43, Hanno: Damit ist nicht alles gelöscht. Bitte im Bericht sauber zwischen Anzeige und Speicherung unterscheiden.
05.10.2026 08:47, Berfin: Ich zähle drei erkannte Referenzfälle, einen übersehenen und drei zusätzliche Marker in unauffälligen Serien.
05.10.2026 08:50, Ottilie: Richtig. Das ist kein repräsentativer Diagnostiktest. Dafür brauchen wir ein anderes Studiendesign.
06.10.2026 11:04, Livia: Wir können zuerst alle Bilder öffnen und die Markierungen danach einblenden. Noch kein ausgelieferter Stand.
06.10.2026 11:09, Hanno: Bitte Version und Testdatum nennen, sobald es soweit ist. Wir ändern bis dahin nichts im Produktivviewer.
07.10.2026 15:20, Berfin: Die Notfallfunktion taucht im Verkaufsprospekt auf. Ist sie in unseren Konten aktiv?
07.10.2026 15:24, Livia: Nein. Eigenes Projekt, kein Teil dieses Labors.
08.10.2026 09:15, Ottilie: Ich liefere noch die medizinische Einschätzung zu den Folgen. Allein meine spätere Unterschrift verhindert keinen übersehenen Befund.
08.10.2026 09:19, Hanno: Danke. Der rechtliche Entwurf muss diese offene Frage sichtbar lassen.
''',xlsx='09_Testdaten_und_Kosten.xlsx',rows=[['Referenz auffällig',4,3,1],['Referenz unauffällig',36,3,33]],costs=[['Testzugänge je Monat',6,240,4],['Netzeinbindung',1,2400,1],['Schulungstermine',2,750,1]],
notes=['B6:B7: Anzahl synthetischer Serien nach vorab festgelegtem Referenzbefund.','C6: richtig markiert; D6: übersehen. C7: zusätzlich markiert; D7: ohne Marker.','40 Serien, davon 4 Referenzauffälligkeiten; keine repräsentative klinische Validierung.','Quelle: 03_Technischer_Fehlertest.pdf, 06_Testlabor_Chat.txt, 07_Klinische_Besprechung.docx.','Kostenquelle: 11_Kostenangebot.pdf. Keine Bestellung, keine fällige Rechnung.'])]

def word(item,out):
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=s.right_margin=Cm(2.25);s.top_margin=s.bottom_margin=Cm(1.9)
    for st in d.styles:
        if st.type==1:st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
    for name,size in [('Title',15),('Heading 1',11)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_after=Pt(8)
    d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Normal'].paragraph_format.line_spacing=1.07
    for n in list(d.styles.element.iter()):
        if n.tag==qn('w:pBdr'):n.getparent().remove(n)
        if n.tag==qn('w:rFonts'):
            for k in list(n.attrib):
                if 'theme' in k.lower():del n.attrib[k]
            n.set(qn('w:ascii'),'Times New Roman');n.set(qn('w:hAnsi'),'Times New Roman')
    d.add_paragraph(item['title'],'Title');d.add_paragraph(item['author'])
    for p in re.split(r'\n\s*\n',item['body']):
        d.add_paragraph(p[3:] if p.startswith('## ') else p,'Heading 1' if p.startswith('## ') else None)
    foot=s.footer.paragraphs[0];foot.alignment=2;foot.add_run('Seite ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    d.core_properties.title=item['title'];d.core_properties.author=item['author'];d.core_properties.created=d.core_properties.modified=datetime(2026,10,9,8)
    d.save(out)

def pdf(item,out,reference):
    normal=Path(os.environ.get('HOCHRISIKO_FONT_REGULAR','/System/Library/Fonts/Supplemental/Times New Roman.ttf'))
    bold=Path(os.environ.get('HOCHRISIKO_FONT_BOLD','/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf'))
    if not normal.exists() or not bold.exists():raise FileNotFoundError('TNR oder explizite Ersatzschrift über HOCHRISIKO_FONT_REGULAR/BOLD angeben.')
    if 'HRSerif' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('HRSerif',str(normal)));pdfmetrics.registerFont(TTFont('HRSerifBold',str(bold)))
    style=ParagraphStyle('Body',fontName='HRSerif',fontSize=11,leading=14,spaceAfter=8)
    heading=ParagraphStyle('Heading',parent=style,fontName='HRSerifBold',spaceBefore=8,spaceAfter=7,keepWithNext=True)
    title=ParagraphStyle('Title',parent=heading,fontSize=15,leading=18,spaceAfter=12)
    story=[Paragraph(html.escape(item['title']),title),Paragraph(html.escape(item['author']),style),Spacer(1,7)]
    for p in re.split(r'\n\s*\n',item['body']):
        h=p.startswith('## ');story.append(Paragraph(html.escape(p[3:] if h else p).replace('\n','<br/>'),heading if h else style))
    def footer(c,doc):
        c.setFont('HRSerif',9);c.drawString(64,30,reference);c.drawRightString(A4[0]-64,30,'Seite '+str(doc.page))
    SimpleDocTemplate(str(out),pagesize=A4,rightMargin=64,leftMargin=64,topMargin=52,bottomMargin=50,title=item['title'],author=item['author']).build(story,onFirstPage=footer,onLaterPages=footer)

def message(item,directory):
    m=EmailMessage(policy=SMTP)
    for h,key in [('From','sender'),('To','to')]:m[h]=tuple(Address(display_name=n,addr_spec=a) for n,a in getaddresses([item[key]]))
    m['Subject']=item['title'];m['Date']=datetime.fromisoformat(item['date']);ident=hashlib.sha256((directory.name+item['file']).encode()).hexdigest()[:24];m['Message-ID']=f'<{ident}@hochrisiko-test.example>'
    m.set_content(item['body']+'\n',charset='utf-8')
    for n in item['attachments']:
        t='pdf' if n.endswith('.pdf') else 'vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        m.add_attachment((directory/n).read_bytes(),maintype='application',subtype=t,filename=n)
    m.set_boundary('hochrisiko-'+ident);(directory/item['file']).write_bytes(m.as_bytes())

def readme(case,d):
    files=sorted([(x['file'],x['title']) for x in case['docs']+case['emails']]+[(case['xlsx'],'Testzählung und Angebotsrechnung mit Formeln'),(case['chatfile'],'Zeitlich zugeordneter Projektchat')])
    rows='\n'.join(f'| [{n}]({n}) | {t} |' for n,t in files)
    text=f'''# 1 {case['title']}

[Alle Testakten](../README.md) · [Repository-Start](../../README.md) · [Download-Index](../../ASSET_INDEX.md)

## 1.1 Arbeitsauftrag und Umfang

{case['intro']}

Bearbeitungsstichtag: 9. Oktober 2026. Einstieg: 01_Auftrag.eml lesen und den vorhandenen Word-Prüfentwurf anhand der übrigen Belege ausarbeiten. Die Akte enthält zwölf eigenständige Arbeitsdateien: vier E-Mails, drei Word-Dateien, drei PDF-Belege, eine Excel-Datei und einen Textchat. Zwei E-Mails enthalten die jeweils genannten Anhänge tatsächlich. Das gibt keine zusätzlichen eigenständigen Belege.

Alle Personen, Einrichtungen, Adressen, Projekte und Vorgänge sind erfunden. Die Kontaktadressen enden auf der reservierten Domain .example. Die Testdaten enthalten keine wirklichen Gerichts- oder Patientenakten. Aussagen einzelner Beteiligter sind Arbeitsmaterial, keine vorweggenommene rechtliche Lösung. Der Arbeitsentwurf ist ausdrücklich nicht fertig.

## 1.2 Originalunterlagen

{WARNING}

| Datei | Inhalt |
| --- | --- |
{rows}

<!-- BEGIN gesamt-pdf-section (autogen) -->

## 1.3 Gesamtausgabe und Archive

{WARNING}

| Format | Download |
| --- | --- |
| Gesamt-PDF | [Gesamtakte](gesamt-pdf/{case['slug']}_gesamt.pdf) |
| Native Originaldateien | [Akten-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/{TAG}/testakte-{case['slug']}.zip) |
| Einzelne Lesefassungen | [Einzel-PDF-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/{TAG}/testakte-{case['slug']}-einzelpdfs.zip) |

<!-- END gesamt-pdf-section (autogen) -->

## 1.4 Passender Arbeitsablauf

[Plugin Hochrisiko-Prüfer](../../ki-verordnung-hochrisiko-pruefer/README.md). Für die reine Einstufung den Artikel-6-Skill wählen; für Vermerk, Anbieterfragen und spätere Fortsetzung den Hauptproblem-Skill verwenden. Die ursprüngliche Recruitingakte bleibt gesondert unverändert erhalten.
'''
    (d/'README.md').write_text(text)
    rubric=f"name: {case['slug']}\nplugin: ki-verordnung-hochrisiko-pruefer\ndescription: {case['title']}\nchecks:\n"
    for i,(n,t) in enumerate(files):rubric+=f'  - id: datei-{i+1}\n    check_type: file_exists\n    description: {t}\n    path: {n}\n'
    rubric+='  - id: dokumentenarbeit\n    check_type: human_review\n    description: Werden Auftrag, Belege, offene Tatsachen und vorhandener Entwurf zu einem ausformulierten Ergebnis zusammengeführt?\n'
    (d/'rubric.yaml').write_text(rubric)

def build(root,messages=False,qa=Path('/tmp/hochrisiko-nachtrag-qa')):
    qa.mkdir(parents=True,exist_ok=True)
    for case in CASES:
        d=root/case['slug'];d.mkdir(parents=True,exist_ok=True)
        if messages:
            for item in case['emails']:message(item,d)
        else:
            for item in case['docs']:
                (word if item['file'].endswith('.docx') else lambda a,b:pdf(a,b,case['reference']))(item,d/item['file'])
            (d/case['chatfile']).write_text(case['chat']);readme(case,d)
    (qa/'cases.json').write_text(json.dumps(CASES,ensure_ascii=False,indent=2)+'\n')
    print('E-Mails erstellt' if messages else 'Dokumente erstellt; Tabellen vor E-Mails erzeugen')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--root',type=Path,default=ROOT/'testakten');p.add_argument('--qa',type=Path,default=Path('/tmp/hochrisiko-nachtrag-qa'));p.add_argument('--messages',action='store_true');a=p.parse_args();build(a.root,a.messages,a.qa)
