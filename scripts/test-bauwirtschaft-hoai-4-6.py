#!/usr/bin/env python3
"""Offline-Regression der Originalakten 4 bis 6, optional kanonische Exporte. Autor: Klotzkette."""
from __future__ import annotations

import argparse
from collections import Counter
import csv
from email import policy
from email.parser import BytesParser
import importlib.util
import io
from pathlib import Path
import re
import sys
import unittest
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageStat
from pypdf import PdfReader
import yaml
from readme_decimal_headings import normalize_decimal_headings
from testakte_disclaimer import NOTICE_BYTES, NOTICE_MARKDOWN, pdf_content_errors

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("hoai456_qa", HERE / "build-bauwirtschaft-hoai-4-6-qa.py")
QA = importlib.util.module_from_spec(spec); spec.loader.exec_module(QA)
ROOT = HERE.parent
ASSETS = None
COUNTS = {4: 17, 5: 17, 6: 18}
FORMATS = {
    4: {".pdf": 7, ".docx": 2, ".xlsx": 1, ".txt": 1, ".eml": 4, ".png": 1, ".csv": 1},
    5: {".pdf": 8, ".docx": 2, ".xlsx": 1, ".txt": 1, ".eml": 3, ".png": 1, ".csv": 1},
    6: {".pdf": 4, ".docx": 3, ".xlsx": 4, ".txt": 1, ".eml": 3, ".png": 1, ".csv": 2},
}


def case(phase): return ROOT / "testakten" / QA.SLUGS[phase]


def text_of(path):
    if path.suffix == ".pdf": return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    if path.suffix == ".docx":
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if path.suffix == ".eml":
        msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        return str(msg) + "\n" + msg.get_body(preferencelist=("plain",)).get_content()
    if path.suffix in (".csv", ".txt"): return path.read_text(encoding="utf-8-sig")
    return ""


class OriginalTests(unittest.TestCase):
    def test_inventory_and_formats(self):
        for phase in COUNTS:
            with self.subTest(phase=phase):
                files = QA.original_files(phase)
                self.assertEqual(len(files), COUNTS[phase])
                self.assertEqual([int(p.name[:2]) for p in files], list(range(1, COUNTS[phase] + 1)))
                self.assertEqual(dict(Counter(p.suffix for p in files)), FORMATS[phase])
                unexpected = [p.name for p in case(phase).iterdir() if p.is_file() and p not in files and p.name not in {"README.md", "rubric.yaml"}]
                self.assertEqual(unexpected, [])
                for path in files:
                    self.assertTrue(path.name.isascii())
                    self.assertGreater(path.stat().st_size, 180)

    def test_readme_and_rubric(self):
        for phase in COUNTS:
            with self.subTest(phase=phase):
                text = (case(phase) / "README.md").read_text()
                self.assertIn("<!-- decimal-headings -->", text)
                self.assertEqual(text, normalize_decimal_headings(text))
                h1 = re.findall(r"^# (.+)$", text, re.M)
                self.assertEqual(len(h1), 1)
                self.assertIn(f"Leistungsphase {phase}", h1[0])
                self.assertRegex(h1[0], r"Bauantrag|Fachkollision|Langtext-LV")
                self.assertRegex(text, re.escape(NOTICE_MARKDOWN) + r"\n\n\| (Fassung|Was)")
                self.assertEqual(text.count("<!-- BEGIN gesamt-pdf-section (autogen) -->"), 1)
                self.assertEqual(text.count("<!-- END gesamt-pdf-section (autogen) -->"), 1)
                self.assertIn("neu erfunden", text)
                self.assertIn("Scheinbar behördliche Schreiben", text)
                self.assertIn("<!-- reserved-example-contacts -->", text)
                for source in QA.original_files(phase): self.assertIn(f"[{source.name}]({source.name})", text)
                rubric = yaml.safe_load((case(phase) / "rubric.yaml").read_text())
                self.assertEqual(rubric["plugin"], "bauwirtschaft")
                self.assertEqual(rubric["author"], "Klotzkette")
                self.assertTrue(any(check["check_type"] == "human_review" for check in rubric["checks"]))

    def test_document_integrity_and_provenance(self):
        for phase in COUNTS:
            for source in QA.original_files(phase):
                with self.subTest(file=source.name, phase=phase):
                    text = text_of(source)
                    self.assertNotIn(chr(167), text)
                    self.assertNotRegex(text, r"Musterlösung|Arbeitsauftrag an die KI|TODO|Lorem ipsum")
                    if source.suffix == ".pdf":
                        reader = PdfReader(source)
                        self.assertGreater(len(text), 180)
                        self.assertEqual(reader.metadata.author, "Klotzkette")
                        self.assertEqual(pdf_content_errors(source.read_bytes()), [])
                        self.assertTrue(all(page.mediabox.width > 500 for page in reader.pages))
                    elif source.suffix == ".docx":
                        doc = Document(source)
                        self.assertEqual(doc.core_properties.author, "Klotzkette")
                        self.assertEqual(doc.core_properties.last_modified_by, "Klotzkette")
                        self.assertGreater(len(text), 1300)
                        self.assertTrue(any(p.style.name == "Title" for p in doc.paragraphs))
                        self.assertIn("gez.", text)
                    elif source.suffix == ".eml":
                        msg = BytesParser(policy=policy.default).parsebytes(source.read_bytes())
                        for field in ("From", "To", "Date", "Subject", "Message-ID"): self.assertTrue(msg[field], field)
                        self.assertIn(".example", msg["From"])
                        self.assertGreater(len(msg.get_body(preferencelist=("plain",)).get_content()), 600)
                    elif source.suffix == ".csv":
                        rows = list(csv.reader(io.StringIO(text), delimiter=";"))
                        self.assertGreaterEqual(len(rows), 5)
                        self.assertGreaterEqual(len(rows[0]), 6)
                        self.assertTrue(all(len(r) == len(rows[0]) for r in rows))
                    elif source.suffix == ".png":
                        with Image.open(source) as im:
                            self.assertGreaterEqual(im.width, 1600)
                            self.assertGreaterEqual(im.height, 1000)
                            self.assertEqual(im.info.get("Author"), "Klotzkette")
                            self.assertGreater(sum(ImageStat.Stat(im.convert("RGB")).var), 150)

    def test_cached_formulas_and_locale(self):
        for phase, name, _, expected, _, _ in QA.SPECS:
            path = case(phase) / name
            with self.subTest(file=name):
                wb = load_workbook(path, data_only=False)
                self.assertEqual(wb.properties.creator, "Klotzkette")
                self.assertEqual(wb.properties.lastModifiedBy, "Klotzkette")
                formulas = QA.formulas(path)
                self.assertGreaterEqual(len(formulas), 3)
                self.assertFalse(any("#REF!" in f for f in formulas.values()))
                if name != "13_Vergabetermine.xlsx":
                    self.assertTrue(any("407" in c.number_format.lower() or "de-de" in c.number_format.lower() for s in wb for row in s for c in row))
                self.assertTrue(all(s.print_area for s in wb))
                wb.close()
                QA.check_values(path, expected)
                cached = load_workbook(path, data_only=True)
                for address in formulas:
                    sheet, cell = address.split("!")
                    self.assertIsNotNone(cached[sheet][cell].value, address)
                cached.close()

    def test_case_specific_evidence(self):
        celle = case(4); hameln = case(5); peine = case(6)
        self.assertIn("48,00 x 18,00", text_of(celle / "03_Lageplan_G01B.pdf"))
        self.assertIn("2,70", text_of(celle / "12_Vermessung.eml"))
        self.assertIn("30.09.2026", text_of(celle / "11_Nachforderung.pdf"))
        self.assertIn("3,15", text_of(hameln / "06_Tragwerksplan_T201B.pdf"))
        self.assertIn("280", text_of(hameln / "07_TGA_Plan_L301D.pdf"))
        self.assertIn("licht 1160", text_of(hameln / "11_Montageplan_M17.pdf"))
        self.assertIn("1,20", text_of(hameln / "03_Entwurfsuebergabe.docx"))
        self.assertIn("Vermaßte Koordinationsdarstellung", text_of(hameln / "04_Grundriss_A101C.pdf"))
        self.assertGreaterEqual(text_of(hameln / "04_Grundriss_A101C.pdf").count("T-01"), 2)
        lv = text_of(peine / "07_LV_Ausbau_ungepreist.docx")
        self.assertEqual(len(re.findall(r"Menge: ", lv)), 16)
        self.assertGreater(len(lv), 8500)
        self.assertNotIn("92.282", lv)
        self.assertNotIn("92282", lv)
        self.assertIn("sechs", text_of(peine / "09_TGA_Schnittstelle.eml"))
        self.assertIn("T-04", text_of(peine / "10_Tuerliste.csv"))
        self.assertIn(";Tür;Raum;", text_of(peine / "10_Tuerliste.csv"))
        self.assertIn("24.12.2026", text_of(peine / "15_Projektkalender.txt"))

    def test_vector_drawing_metadata(self):
        expected = {4: ["03_Lageplan_G01B.pdf", "04_Bauzeichnung_G02B.pdf"], 5: ["04_Grundriss_A101C.pdf", "06_Tragwerksplan_T201B.pdf", "07_TGA_Plan_L301D.pdf", "11_Montageplan_M17.pdf", "16_Fensterdetail_D12.pdf"], 6: ["04_Grundriss_A601C.pdf"]}
        for phase, names in expected.items():
            for name in names:
                with self.subTest(phase=phase, drawing=name):
                    reader = PdfReader(case(phase) / name)
                    self.assertEqual(reader.metadata.subject, "Technische Bauzeichnung")
                    operations = reader.pages[0].get_contents().operations
                    self.assertGreaterEqual(sum(op in (b"m", b"l", b"re", b"c") for _, op in operations), 20)
                    self.assertGreaterEqual(sum(op in (b"S", b"s", b"f", b"f*", b"B", b"B*", b"b", b"b*") for _, op in operations), 10)

    def test_builders_are_scoped(self):
        for name in ("build-bauwirtschaft-hoai-4-6-akten.py", "build-bauwirtschaft-hoai-4-6-workbooks.mjs", "build-bauwirtschaft-hoai-4-6-qa.py"):
            text = (HERE / name).read_text()
            self.assertNotIn("/Users/", text)
            self.assertNotIn("openpyxl.Workbook", text)
            self.assertNotIn("shutil.rmtree", text)
            self.assertNotIn("gesamt-pdf).unlink", text)
        self.assertIn("loadWorkbookRuntime", (HERE / "build-bauwirtschaft-hoai-4-6-workbooks.mjs").read_text())


class ExportTests(unittest.TestCase):
    def test_canonical_assets(self):
        if ASSETS is None: self.skipTest("Optionale kanonische Exporte: --assets DIR")
        for phase, count in COUNTS.items():
            slug = QA.SLUGS[phase]
            originals = QA.original_files(phase)
            for pdf_mode in (False, True):
                suffix = "-einzelpdfs" if pdf_mode else ""
                path = ASSETS / f"testakte-{slug}{suffix}.zip"
                with self.subTest(asset=path.name), zipfile.ZipFile(path) as archive:
                    names = archive.namelist()
                    self.assertEqual(len(names), len(set(names)))
                    self.assertTrue(all("/" not in name and "\\" not in name and name.isascii() for name in names))
                    notice = archive.read("README.txt")
                    self.assertEqual(notice, NOTICE_BYTES)
                    expected = {p.with_suffix(".pdf").name if pdf_mode else p.name for p in originals} | {"README.txt"}
                    combined = f"{slug}_gesamt.pdf"
                    if not pdf_mode:
                        expected.add(combined)
                        self.assertEqual(archive.read(combined), (case(phase) / "gesamt-pdf" / f"{slug}_gesamt.pdf").read_bytes())
                    self.assertEqual(set(names), expected)
                    for source in originals:
                        name = source.with_suffix(".pdf").name if pdf_mode else source.name
                        data = archive.read(name)
                        if name.endswith(".pdf"): self.assertEqual(pdf_content_errors(data), [])
                        if not pdf_mode: self.assertEqual(data, source.read_bytes())
            total = case(phase) / "gesamt-pdf" / f"{slug}_gesamt.pdf"
            data = total.read_bytes()
            self.assertEqual(pdf_content_errors(data), [])
            text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
            for source in originals: self.assertIn(source.name, text)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--assets", type=Path)
    args, extra = parser.parse_known_args()
    ASSETS = args.assets
    unittest.main(argv=[sys.argv[0], *extra], verbosity=2)
