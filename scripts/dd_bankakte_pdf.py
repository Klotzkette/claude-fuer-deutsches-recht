"""Lesbare Druckansicht ausschließlich der Bank-DD-Arbeitsmappe.

Die gespeicherten Excel-Ergebnisse werden gelesen, nicht neu berechnet. Breite
Tabellen erhalten Spaltengruppen mit wiederholten Schlüsseln statt einer
Verkleinerung auf Seitenbreite. Eine Abdeckungsprüfung verhindert Zellverlust.
"""
from __future__ import annotations

from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
import io
from pathlib import Path
import re
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle

FILENAME = '11_Portfolio_und_Kaufpreis.xlsx'
SHEETS = ('Ueberblick', 'Portfolio', 'Buchungen', 'Stichprobe', 'Anleitung')
PAGE = landscape(A4)
WIDTH = PAGE[0] - 72


def display(cell):
    """Angezeigte Zellwerte: Geld centgenau, Prozent wie in der Arbeitsmappe."""
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
        if '%' in cell.number_format:
            return f'{number * 100:.1f}'.replace('.', ',') + '%'
        if '0.00' in cell.number_format:
            number = number.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            return f'{number:,.2f}'.translate(str.maketrans({',': '.', '.': ','}))
        return str(int(number)) if number == number.to_integral_value() else str(number).replace('.', ',')
    return str(value)


def source_cells(path):
    workbook = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    if tuple(workbook.sheetnames) != SHEETS:
        raise ValueError('Die fachlich festgelegten fünf Tabellenblätter weichen ab.')
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                cached = workbook[sheet.title][cell.coordinate]
                if cached.data_type == 'e' or (cell.data_type == 'f' and cached.value is None):
                    raise ValueError(f'Fehlender oder fehlerhafter Formelcache: {sheet.title}!{cell.coordinate}')
    return workbook, {f'{sheet.title}!{cell.coordinate}': display(cell)
                      for sheet in workbook for row in sheet for cell in row if cell.value is not None}


def _fonts():
    candidates = (
        (Path('/System/Library/Fonts/Supplemental'), 'Times New Roman.ttf', 'Times New Roman Bold.ttf'),
        (Path('/usr/share/fonts/truetype/msttcorefonts'), 'Times_New_Roman.ttf', 'Times_New_Roman_Bold.ttf'),
        (Path('/usr/share/fonts/truetype/liberation2'), 'LiberationSerif-Regular.ttf', 'LiberationSerif-Bold.ttf'),
    )
    for directory, regular, bold in candidates:
        if (directory / regular).is_file() and (directory / bold).is_file():
            pdfmetrics.registerFont(TTFont('DD-Bank', str(directory / regular)))
            pdfmetrics.registerFont(TTFont('DD-Bank-Bold', str(directory / bold)))
            return 'Times New Roman' if 'Times' in regular else 'Liberation Serif (Ersatz für Times New Roman)'
    raise RuntimeError('Times New Roman oder Liberation Serif erforderlich.')


def render(path):
    path = Path(path)
    if path.name != FILENAME:
        raise ValueError('Diese Druckansicht gilt ausschließlich für die Bank-DD-Arbeitsmappe.')
    workbook, expected = source_cells(path)
    font_label = _fonts()
    body = ParagraphStyle('DD-Bank-Zelle', fontName='DD-Bank', fontSize=11, leading=14,
                          splitLongWords=True, spaceAfter=0)
    heading = ParagraphStyle('DD-Bank-Kopf', parent=body, fontName='DD-Bank-Bold', textColor=colors.HexColor('#17374B'))
    title = ParagraphStyle('DD-Bank-Titel', parent=heading, fontSize=16, leading=20, spaceAfter=10)
    right = ParagraphStyle('DD-Bank-Zahl', parent=body, alignment=TA_RIGHT)
    covered = set()
    flow = []

    def cell(sheet, coordinate, style=body):
        value = sheet[coordinate]
        key = f'{sheet.title}!{coordinate}'
        if value.value is not None:
            covered.add(key)
        return Paragraph(escape(display(value)), style)

    def prelude(sheet, section):
        if flow:
            flow.append(PageBreak())
        flow.append(Paragraph(escape(section), heading))
        flow.append(Spacer(1, 9))
        for row in sheet.iter_rows(min_row=1, max_row=4):
            for value in row:
                if value.value is not None:
                    flow.append(cell(sheet, value.coordinate, title if value.row == 1 else body))
                    flow.append(Spacer(1, 8))

    def table(sheet, rows, columns, widths, header=True):
        values = []
        for index, row in enumerate(rows):
            rendered = []
            for column in columns:
                source = sheet[f'{column}{row}']
                style = heading if header and index == 0 else right if isinstance(source.value, (int, float)) else body
                rendered.append(cell(sheet, source.coordinate, style))
            values.append(rendered)
        obj = LongTable(values, colWidths=[WIDTH * width / sum(widths) for width in widths],
                        repeatRows=1 if header else 0, hAlign='LEFT', splitByRow=True)
        obj.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LINEBELOW', (0, 0), (-1, 0), 0.6, colors.HexColor('#527487')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F5F7')]),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DCE8EE')),
        ]))
        flow.extend((obj, Spacer(1, 14)))

    sheet = workbook['Ueberblick']
    prelude(sheet, '1. Überblick und Stichtagsüberleitung')
    table(sheet, range(6, 14), ('A', 'B', 'D', 'E'), (240, 130, 220, 180))
    table(sheet, range(16, 21), ('A', 'B', 'D'), (240, 130, 400))
    table(sheet, range(23, 27), ('A', 'B', 'D'), (240, 130, 400))
    for row in (29, 32, 35):
        flow.extend((cell(sheet, f'A{row}'), Spacer(1, 12)))

    groups = (
        ('Portfolio', '2.1 Bestand: Identität, Umfang und Saldo', 'ABCDEFG', (74, 168, 90, 108, 108, 104, 118), 1005),
        ('Portfolio', '2.2 Bestand: Verkäuferpreis, Status und Herkunft', 'AHIJK', (74, 103, 112, 212, 269), 1005),
        ('Buchungen', '3.1 Buchungen: Anfangskapital und neue Belastungen', 'ABCDE', (90, 120, 180, 190, 190), 245),
        ('Buchungen', '3.2 Buchungen: Zahlung und Aufteilung', 'ABFGHI', (90, 120, 140, 140, 140, 140), 245),
        ('Buchungen', '3.3 Buchungen: Endsalden und Buchungstext', 'ABJKLMN', (76, 100, 98, 85, 85, 98, 228), 245),
        ('Stichprobe', '4. Auswahl und Umfang der Detailakten', 'ABCDEF', (74, 140, 196, 135, 95, 130), 25),
        ('Anleitung', '5. Arbeitsweise und Quellen', 'ABCD', (140, 275, 195, 160), 23),
    )
    for name, section, columns, widths, end in groups:
        sheet = workbook[name]
        prelude(sheet, section)
        if name in ('Portfolio', 'Buchungen'):
            flow.extend((Paragraph('Fortsetzung derselben Datensätze in Spaltengruppen; '
                                   'Darlehensnummer' + (' und Buchungsdatum' if name == 'Buchungen' else '') +
                                   ' werden in jeder Gruppe wiederholt. Geldbeträge in EUR.', body), Spacer(1, 10)))
        table(sheet, range(5, end + 1), tuple(columns), widths)
    if covered != set(expected):
        raise ValueError(f'Nicht abgebildete Zellen: {sorted(set(expected) - covered)}')

    output = io.BytesIO()
    def footer(canvas, document):
        canvas.setFont('DD-Bank', 9)
        canvas.drawString(36, 21, f'Bank-DD | vollständige Druckansicht | {font_label} 11 pt')
        canvas.drawRightString(PAGE[0] - 36, 21, f'Seite {document.page}')
    document = SimpleDocTemplate(output, pagesize=PAGE, leftMargin=36, rightMargin=36,
                                 topMargin=32, bottomMargin=40, invariant=1,
                                 title='Bank-DD: vollständige Arbeitsmappe in lesbaren Spaltengruppen',
                                 author='Kanzleiakte')
    document.build(flow, onFirstPage=footer, onLaterPages=footer)
    data = output.getvalue()
    assert_preserved(path, data)
    return data


def assert_preserved(path, data):
    """Alle nichtleeren Cache-Zellen und ihre Geldbeträge müssen gedruckt sein."""
    _, expected = source_cells(path)
    text = '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(data)).pages)
    compact = re.sub(r'\s+', '', text)
    missing = [coordinate for coordinate, value in expected.items()
               if re.sub(r'\s+', '', value) not in compact]
    if missing:
        raise ValueError(f'Zellwerte fehlen im PDF-Text: {missing[:20]}')
    amounts = Counter(value for value in expected.values() if re.fullmatch(r'-?[\d.]+,\d{2}', value))
    printed = Counter(re.findall(r'(?<![\w.,])-?[\d.]+,\d{2}(?![\d%])', text))
    if amounts - printed:
        raise ValueError(f'Geldbeträge fehlen in der Druckansicht: {amounts - printed}')
    if '\ufffd' in text or '\u25a0' in text:
        raise ValueError('Nicht druckbare Zeichen in der Arbeitsmappe.')
