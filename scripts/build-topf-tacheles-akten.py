#!/usr/bin/env python3
"""Erzeugt ausschließlich acht native Aktenstücke der einfachen Gründungsakte.

Aufruf mit dem gebündelten Codex-Python. Die sieben JPG-Datenkarten und die
Auslieferungspakete werden gesondert erstellt. Es werden keine PDFs erzeugt.
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "scripts/data/topf-tacheles/fakten.json"
F = json.loads(DATA.read_text(encoding="utf-8"))
OUT = ROOT / "testakten" / F["slug"]
OUT.mkdir(parents=True, exist_ok=True)


def euro(amount: int) -> str:
    return f"{amount:,}".replace(",", ".") + ",00 EUR"


def p(doc, text, bold=False):
    after_table = len(doc._element.body) >= 2 and doc._element.body[-2].tag == qn("w:tbl")
    para = doc.add_paragraph()
    if after_table:
        para.paragraph_format.space_before = Pt(6)
    run = para.add_run(text)
    run.bold = bold
    return para


def heading(doc, text):
    return doc.add_paragraph(text, style="Heading 1")


def clause(doc, number, text):
    para = doc.add_paragraph()
    para.add_run(number + "  ").bold = True
    para.add_run(text)
    return para


def document(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.35), Cm(2.2)
    sec.header_distance, sec.footer_distance = Cm(0.9), Cm(0.8)
    styles = doc.styles
    for st in styles:
        for border in list(st._element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
        st = styles[name]
        st.font.name = "Times New Roman"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(11)
        fonts = st._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if "theme" in attr.lower():
                del fonts.attrib[attr]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + script), "Times New Roman")
    normal = styles["Normal"].paragraph_format
    normal.line_spacing = 1.10
    normal.space_after = Pt(6)
    normal.widow_control = True
    styles["Title"].font.size = Pt(17)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(8)
    styles["Subtitle"].paragraph_format.space_after = Pt(10)
    styles["Subtitle"].font.italic = False
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(11)
    styles["Heading 1"].paragraph_format.space_after = Pt(11)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.text = "Topf & Tacheles GmbH · Gründungsvorbereitung"
    header.runs[0].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("ENTWURF · Stand 30.09.2026 · Seite ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
    doc.add_paragraph(title, style="Title")
    doc.add_paragraph(subtitle, style="Subtitle")
    doc.core_properties.author = "Gründungskreis Topf & Tacheles"
    doc.core_properties.title = title
    doc.core_properties.subject = "ENTWURF · Stand 30.09.2026"
    doc.core_properties.created = datetime(2026, 9, 30, 7, 0, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 9, 30, 7, 0, tzinfo=timezone.utc)
    return doc


def table(doc, headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    for col, width in zip(tbl.columns, widths):
        col.width = Cm(width)
    for cell, width in zip(tbl.rows[0].cells, widths):
        cell.width = Cm(width)
    for i, label in enumerate(headers):
        tbl.rows[0].cells[i].text = label
    trpr = tbl.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    trpr.append(repeat)
    for values in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            cells[i].width = Cm(widths[i])
    for ri, row in enumerate(tbl.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge, value in (("top", "70"), ("bottom", "70"), ("left", "80"), ("right", "80")):
                item = OxmlElement("w:" + edge)
                item.set(qn("w:w"), value)
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            tcpr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "bottom", "left", "right"):
                item = OxmlElement("w:" + edge)
                item.set(qn("w:val"), "single")
                item.set(qn("w:sz"), "4")
                item.set(qn("w:color"), "D9D9D9")
                borders.append(item)
            tcpr.append(borders)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EAEAEA")
                tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                if ci == len(headers) - 1:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in para.runs:
                    run.bold = ri == 0
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)
    return tbl


def mail(filename, sender, to, subject, time, body):
    msg = EmailMessage(policy=SMTP)
    msg["From"] = sender
    msg["To"] = to
    msg["Date"] = f"Wed, 30 Sep 2026 {time}:00 +0200"
    msg["Subject"] = subject
    msg["Message-ID"] = f"<{filename[:2]}.20260930@topf-tacheles.example>"
    msg.set_content(body.strip() + "\n", charset="utf-8")
    (OUT / filename).write_bytes(msg.as_bytes())


mail("01_Auftrag_Frieda.eml", "Frieda <frieda@topf-tacheles.example>",
     "Gründungsberatung <beratung@kanzlei.example>", "Topf & Tacheles – bitte zwei kurze Entwürfe fertig machen", "08:45", """
Guten Morgen,

wir sind sieben Leute und wollen in Berlin die Topf & Tacheles GmbH gründen. Wir vermieten Zimmerpflanzen und übernehmen Lieferung, Gießen und Pflege. Wir brauchen eine einfache Satzung und eine kurze Gesellschaftervereinbarung.

Bitte füllen Sie die Namen, Geburtsdaten und Wohnorte in unseren beiden Word-Entwürfen aus den sieben JPG-Datenkarten aus. P01 bis P07 müssen bei den richtigen Anteilen bleiben. Lio heißt auf seiner Karte Leopold; bitte seinen vollständigen Namen verwenden.

Das Stammkapital soll 25.000 Euro betragen. Die Aufteilung 30 / 25 / 15 / 10 / 10 / 5 / 5 Prozent ist unser Plan. Wir wollen bar einzahlen; Pflanzen, Arbeitszeit und Lastenrad sollen darauf nicht angerechnet werden.

Bitte prüfen Sie besonders die Stimmen und Mehrheiten. Lio meint, seine 25 Prozent seien eine Sperre. Bruno denkt, jeder hätte eine Stimme. Ich möchte den Alltag organisieren können, ohne für jeden Blumentopf sieben Zusagen einzusammeln. Bei größeren Käufen und Krediten wollen wir gemeinsam entscheiden. Beträge und Mehrheiten sind noch offen. Bitte offene Entscheidungen deutlich stehen lassen und uns nur dazu kurze Vorschläge machen.

Ich wäre als Geschäftsführerin bereit, bestellt ist bisher niemand. Bitte schicken Sie uns zunächst die ergänzten Entwürfe zur gemeinsamen Durchsicht. Die Satzung ist nicht beurkundet. Ein Gesellschaftskonto und Einzahlungen gibt es noch nicht.

Vielen Dank
Frieda
""")

(OUT / "02_Kurzer_Gruenderchat.txt").write_text("""Gruppe „Topf & Tacheles“
Auszug vom 28. und 29. September 2026 · zusammengestellt am 30. September 2026

28.09.2026, 18:02 – Frieda (P01): Ich würde den Betrieb führen. Pflanzen und Pflege kenne ich; für die Buchhaltung brauche ich Nepomuk.
28.09.2026, 18:05 – Lio (P02): Touren und Verkauf übernehme ich gern. Bei meiner früheren GmbH gab es eine Satzung. Ich habe uns einen kurzen Entwurf angefangen.
28.09.2026, 18:08 – Hatice (P03): Für mich ist das die erste Gründung. Mit 15 Prozent und den 3.750 Euro bin ich einverstanden.
28.09.2026, 18:10 – Nepomuk (P04): Der Anteilsplan kommt auf 25.000 Euro. Bitte nicht mit dem Geld verwechseln, das später für Miete und Pflanzen gebraucht wird.
28.09.2026, 18:14 – Amina (P05): Ich kümmere mich um Gestaltung und Internetseite. Meine Arbeit ist aber nicht meine Einzahlung, richtig? Die 2.500 Euro plane ich extra ein.
28.09.2026, 18:16 – Frieda (P01): Ja, wir bleiben bei Geld. Meine vorhandenen Pflanzen gehen nicht automatisch in die GmbH. Über eine Nutzung reden wir später.
28.09.2026, 18:19 – Bruno (P06): Ich zahle 1.250 Euro und helfe bei Lieferungen. Wenn wir gleichberechtigt sind, hat dann jeder von uns eine Stimme?
28.09.2026, 18:23 – Lio (P02): In meinem Text steht eine Stimme je Euro. Für große Sachen wollte ich drei Viertel. Mit meinen 25 Prozent könnte ohne mich dann nichts durchgehen, dachte ich.
28.09.2026, 18:27 – Hatice (P03): Das sollten wir prüfen lassen. Ich habe weder eine besondere Sperre noch Kopfstimmen zugesagt.
29.09.2026, 09:04 – Zoe (P07): Ich mache Kundenbefragungen und soziale Medien. Ich möchte früh genug von wichtigen Entscheidungen erfahren, auch mit fünf Prozent.
29.09.2026, 09:08 – Frieda (P01): Alltag gern bei mir. Einen teuren Lieferwagen oder einen großen Kredit möchte ich nicht allein entscheiden. Wo die Grenze liegt, weiß ich noch nicht.
29.09.2026, 09:12 – Nepomuk (P04): Bitte auch klären, ob bei einer Mehrheit alle Anteile zählen oder nur die Leute, die abstimmen. Ich möchte das nachlesen können.
29.09.2026, 09:16 – Amina (P05): Die private Vereinbarung bitte einfach halten. Aufgaben und Bezahlung besprechen wir noch. Nicht jetzt schon feste Wochenstunden eintragen.
29.09.2026, 09:20 – Frieda (P01): Ich gebe die zwei Entwürfe und unsere Datenkarten zur Bearbeitung weiter. Erst gemeinsam lesen, dann zum Notariat. Noch nichts unterschreiben oder einzahlen.
""", encoding="utf-8")

doc = document("Vorüberlegungen und Anteilsplan", "Arbeitsnotiz von Frieda und Lio · ENTWURF · Stand 30.09.2026")
p(doc, "Wir wollen zu siebt einen überschaubaren Pflanzenverleih aufbauen. Diese Notiz hält unseren bisherigen Plan fest. Die persönlichen Angaben stehen auf den sieben Datenkarten; die Kürzel bleiben in allen Unterlagen gleich.")
heading(doc, "1 Geschäft und Stand")
p(doc, "Die Topf & Tacheles GmbH soll ihren Sitz in Berlin haben. Wir wollen robuste Zimmerpflanzen an Berliner Büros und Wohngemeinschaften vermieten, liefern und regelmäßig gießen und pflegen. Pflanzen und Pflanzgefäße sollen ergänzend verkauft werden. Der Firmenname ist noch nicht verbindlich geprüft.")
p(doc, "Als Geschäftsanschrift ist Fiktive Topfgasse 12, 10405 Berlin vorgesehen. Es gab noch keine Beurkundung und keine Handelsregistereintragung. Ein Gesellschaftskonto und Einzahlungen gibt es noch nicht.")
heading(doc, "2 Personen und geplanter Kapitaleinsatz")
roles = ["Betrieb und Pflanzenpflege", "Touren und Verkauf", "Pflanzen und Pflege", "Buchhaltung", "Gestaltung und Internetauftritt", "Gelegentliche Lieferhilfe", "Kundenbefragungen und soziale Medien"]
rows = []
for person, role in zip(F["personen"], roles):
    name = person["vorname"] + (" (Lio)" if person["id"] == "P02" else "")
    rows.append([person["id"], name, role, euro(person["nennbetrag"]), f'{person["quote"]} %'])
table(doc, ["Karte", "Vorname", "Vorgesehene Aufgabe", "Nennbetrag", "Anteil"], rows, [1.5, 2.6, 6.0, 3.3, 1.35])
p(doc, "Zusammen übernehmen wir sieben Geschäftsanteile mit Nennbeträgen von 25.000,00 EUR und damit 100 Prozent. Jede Person übernimmt genau den ihrer Karte zugeordneten Anteil. Die Aufgaben sind eine Arbeitsaufteilung im Gespräch, keine vereinbarten Arbeitsverträge.")
heading(doc, "3 Erfahrung im Team")
p(doc, "Frieda führt bisher einen kleinen Einzelbetrieb; Lio hat schon eine GmbH mitgegründet. Hatice kennt den Pflanzenhandel, Nepomuk die Buchhaltung und Amina die Selbstständigkeit. Bruno und Zoe gründen zum ersten Mal. Wir brauchen deshalb verständliche Texte, die auch ohne Vorerfahrung lesbar sind.")
doc.add_page_break()
heading(doc, "4 Geld und Mitarbeit")
p(doc, "Wir wollen eine Bargründung. Arbeitsleistungen, Pflanzen und ein vorhandenes Lastenrad werden nicht auf die Einlagen angerechnet. Falls die GmbH solche Sachen später nutzt, brauchen wir dazu eine getrennte Vereinbarung. Auch Vergütung und Zeitumfang der Mitarbeit sind noch offen.")
p(doc, "Die Aufteilung des Stammkapitals steht als gemeinsamer Plan. Wir haben noch keinen Zahlungstermin verabredet. Das Geld für den späteren laufenden Betrieb müssen wir daneben planen; zusätzliche Zahlungsverpflichtungen wollen wir mit diesen Entwürfen noch nicht eingehen.")
heading(doc, "5 Stimmen und Entscheidungen")
p(doc, "Lios Satzungsentwurf zählt eine Stimme je Euro des Geschäftsanteils. Bruno ging bislang von einer Stimme je Person aus. Lio erwartet bei einer Dreiviertelmehrheit mit 25 Prozent eine Sperre. Frieda und Hatice haben dazu nichts zugesagt. Wir wollen vor der Einigung verstehen, was bei allen Anwesenden, bei Abwesenheit und bei Enthaltungen herauskommt.")
p(doc, "Frieda soll nach dem bisherigen Vorschlag die Geschäftsführung übernehmen. Im Alltag soll sie handlungsfähig sein. Größere Anschaffungen und Kredite sollen eine gemeinsame Entscheidung benötigen. Die Geldgrenzen, die erforderliche Mehrheit und deren Bezugsbasis fehlen noch. Auch die Vertretungsregelung ist noch zu besprechen.")
heading(doc, "6 Auftrag für die nächste Fassung")
p(doc, "In die Satzung und die Gesellschaftervereinbarung sollen die vollständigen Namen und die jeweiligen Personendaten aus P01 bis P07 übernommen werden. Die zugeordneten Geschäftsanteile dürfen dabei nicht verrutschen. Die offenen Entscheidungsfelder sollen sichtbar bleiben, bis wir sie gemeinsam entschieden haben.")
p(doc, "Wir möchten anschließend zwei kurze, zusammenpassende Entwürfe lesen. Die Satzung soll die Regeln der Gesellschaft enthalten. Die zusätzliche Vereinbarung soll unsere Zusammenarbeit festhalten. Bevor etwas beurkundet wird, wollen alle sieben die ausgefüllten Texte sehen.")
doc.save(OUT / "03_Vorueberlegungen_und_Anteilsplan.docx")

doc = document("Satzung zur Gründung", "Topf & Tacheles GmbH · ENTWURF · Stand 30.09.2026")
p(doc, "Die folgenden Regelungen sind ein noch nicht vereinbarter Entwurf. Namen und Personendaten sind aus den zugeordneten Datenkarten einzutragen. Als OFFEN bezeichnete Felder erfordern eine gemeinsame Entscheidung; diese Fassung ist nicht zur Beurkundung freigegeben.")
heading(doc, "1 Firma Sitz und Unternehmensgegenstand")
clause(doc, "1.1", "Die Firma der Gesellschaft lautet Topf & Tacheles GmbH. Die Gesellschaft hat ihren Sitz in Berlin.")
clause(doc, "1.2", "Gegenstand des Unternehmens ist die Vermietung von Zimmerpflanzen einschließlich Lieferung sowie Gieß- und Pflegeservice für Büros und Wohngemeinschaften und der ergänzende Verkauf von Pflanzen und Pflanzgefäßen.")
clause(doc, "1.3", "Die Gesellschaft wird auf unbestimmte Zeit errichtet. Das Geschäftsjahr ist das Kalenderjahr; das erste Geschäftsjahr endet am 31. Dezember des Jahres der Eintragung in das Handelsregister.")
heading(doc, "2 Stammkapital und Geschäftsanteile")
clause(doc, "2.1", "Das Stammkapital beträgt 25.000,00 EUR. Die folgenden sieben Personen übernehmen jeweils einen Geschäftsanteil mit dem zugeordneten Nennbetrag. Die Kürzel P01 bis P07 dienen nur der Zuordnung zu den Datenkarten.")
table(doc, ["Anteil Nr", "Übernehmende Person", "Nennbetrag"],
      [[str(i), f'{x["id"]}: [vollständiger Name aus Karte {x["id"]}]', euro(x["nennbetrag"])] for i, x in enumerate(F["personen"], 1)], [1.8, 9.0, 3.55])
p(doc, "Die weiteren Angaben zu den übernehmenden Personen lauten für jede Karte: [P01 bis P07 jeweils: Geburtsdatum und Wohnort aus der entsprechenden Datenkarte]. Diese Angaben sind vor der Beurkundung personengenau zu ergänzen.")
clause(doc, "2.2", "Die Einlagen sind in Geld zu leisten. Auf jeden Geschäftsanteil ist vor der Anmeldung zur Eintragung mindestens ein Viertel des Nennbetrags einzuzahlen; insgesamt müssen mindestens 12.500,00 EUR eingezahlt sein. Der weitere Einzahlungsplan lautet: [OFFEN: Höhe und Fälligkeit der ersten Zahlung je Person sowie Fälligkeit der Restbeträge]. Die genannten gesetzlichen Mindestzahlungen bleiben verbindlich.")
doc.add_page_break()
heading(doc, "3 Geschäftsführung und Vertretung")
clause(doc, "3.1", "Die Gesellschaft hat eine oder mehrere geschäftsführende Personen. Ihre Bestellung und Abberufung erfolgen durch Beschluss der Gesellschafterversammlung. Als erste Geschäftsführerin ist die Person P01 vorgesehen; ihre Bestellung ist einem gesonderten Beschluss vorbehalten.")
clause(doc, "3.2", "Ist nur eine geschäftsführende Person bestellt, vertritt sie die Gesellschaft allein. Sind mehrere bestellt, vertreten sie die Gesellschaft gemeinschaftlich. Die Gesellschafterversammlung kann einzelnen geschäftsführenden Personen Einzelvertretungsbefugnis erteilen. Diese vorgeschlagene Vertretungsregelung soll [OFFEN: beibehalten oder durch die gemeinsam festgelegte Vertretungsregel ersetzt werden].")
clause(doc, "3.3", "Die Geschäftsführung führt die gewöhnlichen Geschäfte der Gesellschaft eigenverantwortlich im Rahmen des Gesetzes, dieser Satzung und der Gesellschafterbeschlüsse. Sie hat vor einer Anschaffung mit einer Gesamtverpflichtung von mehr als [OFFEN: Betrag in EUR] sowie vor Aufnahme eines Kredits von mehr als [OFFEN: Betrag in EUR] die Zustimmung der Gesellschafterversammlung nach Ziffer 5.2 einzuholen.")
clause(doc, "3.4", "Die Zustimmungspflichten beschränken die Geschäftsführung im Innenverhältnis. Sie beschränken die gesetzliche Vertretungsmacht gegenüber Dritten nicht.")
heading(doc, "4 Gesellschafterversammlung und Information")
clause(doc, "4.1", "Die Geschäftsführung beruft die Gesellschafterversammlung durch eingeschriebenen Brief mit einer Frist von mindestens einer Woche ein. Die Einladung nennt den Ort, die Zeit und die Tagesordnung. Die Versammlung findet am Sitz der Gesellschaft statt, sofern nicht alle Gesellschafterinnen und Gesellschafter einem anderen Ort zustimmen.")
clause(doc, "4.2", "Eine Versammlung per Telefon oder Video sowie eine Beschlussfassung ohne Versammlung sind unter den jeweiligen gesetzlichen Voraussetzungen zulässig. Jeder Person muss eine Teilnahme beziehungsweise Stimmabgabe im gesetzlich vorgesehenen Verfahren ermöglicht werden.")
clause(doc, "4.3", "Über die Beschlüsse wird eine Niederschrift mit Gegenstand und Abstimmungsergebnis angefertigt und allen Gesellschafterinnen und Gesellschaftern übermittelt. Die gesetzlich erforderliche notarielle Beurkundung einzelner Beschlüsse wird dadurch nicht ersetzt.")
clause(doc, "4.4", "Jede Gesellschafterin und jeder Gesellschafter hat die gesetzlichen Auskunfts- und Einsichtsrechte. Die Geschäftsführung stellt die zur Vorbereitung eines angekündigten Beschlusses erforderlichen Unterlagen rechtzeitig zur Verfügung.")
doc.add_page_break()
heading(doc, "5 Stimmrechte und Mehrheiten")
clause(doc, "5.1", "Jeder Euro des Nennbetrags eines Geschäftsanteils gewährt eine Stimme. Gewöhnliche Beschlüsse werden mit einfacher Mehrheit der abgegebenen gültigen Stimmen gefasst, soweit das Gesetz oder diese Satzung keine strengere Anforderung vorsieht. Ein Antrag ist angenommen, wenn die Ja-Stimmen die Nein-Stimmen überwiegen. Enthaltungen und ungültige Stimmen werden nicht mitgezählt; bei Stimmengleichheit ist der Antrag abgelehnt. Gesetzliche Stimmverbote bleiben unberührt.")
clause(doc, "5.2", "Die Zustimmung zu den in Ziffer 3.3 genannten größeren Anschaffungen und Krediten erfordert [OFFEN: mindestens oder mehr als welcher Prozentsatz] der [OFFEN: Bezugsbasis der Stimmen ausdrücklich festlegen]. Bei dieser Bezugsbasis werden Enthaltungen und Abwesenheiten wie folgt behandelt: [OFFEN: Behandlung festlegen]. Diese Felder müssen vor der abschließenden Fassung gemeinsam ausgefüllt werden.")
clause(doc, "5.3", "Eine Änderung der Satzung bedarf mindestens einer Mehrheit von drei Vierteln der abgegebenen Stimmen. Der Beschluss muss notariell beurkundet werden. Eine Vermehrung der den Gesellschafterinnen und Gesellschaftern nach der Satzung obliegenden Leistungen bedarf außerdem der Zustimmung sämtlicher betroffener Personen. Zusätzliche Anforderungen sollen lauten: [OFFEN: keine zusätzlichen Anforderungen oder genaue zusätzliche Mehrheit, Bezugsbasis und gegebenenfalls Zustimmungserfordernisse]. Gesetzliche Mindestmehrheiten, Formvorschriften und Zustimmungserfordernisse bleiben in jedem Fall unberührt.")
clause(doc, "5.4", "Für sonstige gesetzlich besonders geregelte Beschlüsse gelten die gesetzlichen Anforderungen. Ein besonderes Vetorecht einer einzelnen Person wird mit diesem Entwurf nicht festgelegt. Ob ein solches Recht gewünscht ist, bleibt zur gemeinsamen Entscheidung offen.")
heading(doc, "6 Jahresabschluss und Ergebnis")
clause(doc, "6.1", "Die Geschäftsführung stellt den Jahresabschluss innerhalb der gesetzlichen Fristen auf und legt ihn der Gesellschafterversammlung zur Feststellung vor. Die Gesellschafterversammlung beschließt über die Verwendung des Ergebnisses.")
clause(doc, "6.2", "Soweit ein ausschüttungsfähiger Gewinn zur Verteilung beschlossen wird, erfolgt die Verteilung im Verhältnis der Nennbeträge der Geschäftsanteile. Die gesetzlichen Vorschriften zur Erhaltung des Stammkapitals sind einzuhalten.")
heading(doc, "7 Gründungsaufwand und Ergänzung")
clause(doc, "7.1", "Die Gesellschaft trägt den Aufwand für die notarielle Gründung, die Handelsregistereintragung und die Gründungsberatung bis zu einem Gesamtbetrag von [OFFEN: Höchstbetrag in EUR]. Darüber hinausgehende Gründungskosten tragen die Gründenden im Verhältnis ihrer Nennbeträge. Der Höchstbetrag ist vor der Beurkundung festzulegen.")
clause(doc, "7.2", "Im Übrigen gelten die gesetzlichen Vorschriften. Diese Entwurfsfassung enthält keine Unterschriften und dokumentiert noch keinen Abschluss des Gesellschaftsvertrags.")
doc.save(OUT / "04_Satzung_Entwurf.docx")

doc = document("Gesellschaftervereinbarung", "Topf & Tacheles GmbH · ENTWURF · Stand 30.09.2026")
p(doc, "Diese Vereinbarung soll die Zusammenarbeit der sieben Gründenden ergänzend zur Satzung regeln. Sie ist noch nicht vereinbart. Die mit OFFEN bezeichneten Entscheidungen bleiben bis zur gemeinsamen Einigung offen.")
heading(doc, "1 Beteiligte und Zweck")
p(doc, "Die folgenden Personen schließen diese Vereinbarung. Die Angaben in eckigen Klammern sind aus der jeweils bezeichneten Datenkarte zu übernehmen.")
for person in F["personen"]:
    p(doc, f'{person["id"]}: [vollständiger Name], geboren am [Geburtsdatum], wohnhaft in [Wohnort], jeweils aus Datenkarte {person["id"]}.')
p(doc, "Die Beteiligten bereiten die Bargründung der Topf & Tacheles GmbH mit Sitz in Berlin vor. Die Geschäftsanteile sollen P01 zu 30 Prozent, P02 zu 25 Prozent, P03 zu 15 Prozent, P04 und P05 zu jeweils 10 Prozent sowie P06 und P07 zu jeweils 5 Prozent zustehen. Maßgeblich für die Übernahme sind die Nennbeträge in der notariell zu errichtenden Satzung.")
heading(doc, "2 Zusammenarbeit und Mitarbeit")
clause(doc, "2.1", "Die Beteiligten informieren einander rechtzeitig über Umstände, die die gemeinsame Gründung oder den Betrieb wesentlich betreffen. Sie behandeln einander fair und besprechen Meinungsverschiedenheiten zunächst unmittelbar miteinander.")
clause(doc, "2.2", "Die Aufgaben sollen wie folgt verteilt werden: P01 betreut den Betrieb und die Pflanzenpflege, P02 die Touren und den Verkauf, P03 die Pflanzen und Pflege, P04 die Buchhaltung, P05 die Gestaltung und den Internetauftritt, P06 gelegentliche Lieferhilfe und P07 Kundenbefragungen und soziale Medien. Der verbindliche Umfang der Mitarbeit lautet: [OFFEN: vereinbarter Umfang je Person oder Verweis auf später gesondert abzuschließende Verträge].")
clause(doc, "2.3", "Für vergütete Tätigkeiten sollen vor ihrem Beginn gesonderte Vereinbarungen mit der Gesellschaft geschlossen werden. Die dafür vorgesehenen Vergütungsgrundsätze lauten: [OFFEN: gemeinsame Vorgaben zur Vergütung]. Dieser Entwurf enthält noch keine Zusage einer bestimmten Vergütung und ersetzt keinen Arbeits- oder Dienstvertrag.")
doc.add_page_break()
heading(doc, "3 Einlagen und zusätzliche Mittel")
clause(doc, "3.1", "Die Einlagen werden ausschließlich in Geld nach der Satzung und dem festzulegenden Zahlungsplan geleistet. Arbeitsleistungen, Pflanzen und das vorhandene Lastenrad werden nicht auf das Stammkapital angerechnet. Eine Nutzung solcher Gegenstände durch die Gesellschaft bedarf einer gesonderten Vereinbarung.")
clause(doc, "3.2", "Mit dieser Vereinbarung werden keine zusätzlichen Einzahlungen, Gesellschafterdarlehen, Bürgschaften oder Nachschüsse zugesagt. Benötigt die Gesellschaft weitere Mittel, werden die Beteiligten darüber gesondert beraten; eine Verpflichtung setzt eine wirksame gesonderte Vereinbarung voraus.")
heading(doc, "4 Gemeinsame Entscheidungen")
clause(doc, "4.1", "Die Beteiligten wirken darauf hin, dass die Geschäftsführung vor Beschlüssen über größere Anschaffungen oder Kredite allen Beteiligten den Zweck, den Betrag und die Auswirkungen auf die verfügbaren Mittel in verständlicher Form mitteilt. Die dafür gewünschte Vorlaufzeit beträgt [OFFEN: Zahl der Tage]. Gesetzliche Einladungsfristen bleiben unberührt.")
clause(doc, "4.2", "Für die Satzung sollen die Beteiligten folgende gemeinsame Entscheidung zu Stimmgewicht, Mehrheiten und besonderen Zustimmungserfordernissen festhalten: [OFFEN: abgestimmte Regelung nach Prüfung der unterschiedlichen Vorstellungen]. Bis zur Einigung ist weder ein persönliches Vetorecht für P02 noch eine Stimme je Person zugesagt.")
clause(doc, "4.3", "Diese Vereinbarung wirkt nur als Vertrag zwischen ihren Beteiligten. Sie ändert die Satzung und die Wirksamkeitsvoraussetzungen von Gesellschafterbeschlüssen nicht. Soweit eine Regel durch Änderung der Satzung umgesetzt werden soll, sind die gesetzlichen Form- und Mehrheitserfordernisse sowie die erforderliche Registereintragung zu beachten. Interne Vorgaben für die Geschäftsführung sind erforderlichenfalls durch einen wirksamen Gesellschafterbeschluss umzusetzen.")
heading(doc, "5 Informationen und Vertraulichkeit")
clause(doc, "5.1", "Die Beteiligten geben vertrauliche Geschäfts- und Kundendaten nur weiter, soweit dies für den Betrieb, die berufliche Beratung oder zur Erfüllung gesetzlicher Pflichten erforderlich ist. Ihre gesetzlichen Informationsrechte gegenüber der Gesellschaft bleiben unberührt.")
heading(doc, "6 Streit und Abschluss")
clause(doc, "6.1", "Bei einem Streit erläutern die Beteiligten ihre Anliegen zunächst in einem gemeinsamen Gespräch. Führt dieses zu keiner Einigung, soll [OFFEN: gewünschtes weiteres Vorgehen] gelten. Der Zugang zu den zuständigen Gerichten bleibt unberührt.")
clause(doc, "6.2", "Diese Vereinbarung kommt erst zustande, wenn alle sieben Beteiligten der vervollständigten Fassung wirksam zugestimmt haben. Änderungen bedürfen der Zustimmung aller Beteiligten und sollen schriftlich festgehalten werden. Gesetzliche Formvorschriften bleiben unberührt. Dieser Entwurf begründet weder eine Pflicht zur Übertragung von Geschäftsanteilen noch eine Kauf- oder Verkaufsoption.")
p(doc, "Ort und Datum des beabsichtigten Abschlusses bleiben offen. Unterschriften sind in dieser Entwurfsfassung nicht enthalten.")
doc.save(OUT / "05_Gesellschaftervereinbarung_Entwurf.docx")

mail("06_Mehrheiten_Rueckfrage.eml", "Lio <lio@topf-tacheles.example>",
     "Gründungsberatung <beratung@kanzlei.example>", "Mehrheiten: Reichen meine 25 Prozent zum Stoppen?", "10:05", """
Guten Tag,

ich möchte nicht, dass die anderen einen großen Kredit ohne mich beschließen. Ich dachte, meine 25 Prozent reichen bei drei Vierteln zum Stoppen. Gilt das auch, wenn alle anderen dafür sind? Und was passiert, wenn jemand fehlt oder sich enthält?

Bruno hat mir noch geschrieben: „Ich dachte, jeder von uns hat eine Stimme. Sonst können Frieda und du ja viel mehr entscheiden als wir anderen.“ Er möchte wissen, ob gleichberechtigt automatisch Kopfstimmen bedeutet.

Frieda will den Alltag allein erledigen können. Hatice möchte aber nicht, dass wir jetzt einfach ein persönliches Vetorecht für mich einbauen. Dazu haben wir noch keine Einigung. Bitte die Stellen in beiden Entwürfen sichtbar offen lassen und die Varianten kurz erklären.

Die Betragsgrenze für Anschaffungen und Kredite haben wir ebenfalls noch nicht festgelegt.

Viele Grüße
Lio (P02)
""")

mail("07_Notariat_Terminvorbereitung.eml", "Notariatsbüro <vorbereitung@notariat.example>",
     "Frieda <frieda@topf-tacheles.example>", "Vorbereitung der Gründung der Topf & Tacheles GmbH", "11:30", """
Guten Tag,

für die Vorbereitung benötigen wir die abgestimmte Satzung und die vollständigen Angaben aller sieben Gründenden. Bitte teilen Sie außerdem mit, wer zur Geschäftsführung bestellt werden soll, welche Vertretungsbefugnis vorgesehen ist und welche Geschäftsanschrift angemeldet werden soll.

Bitte klären Sie zuvor die noch offenen Mehrheitsregeln, den Zahlungsplan und den Höchstbetrag des Gründungsaufwands. Der gewünschte Firmenname muss noch geprüft werden. Die vorliegenden Datenkarten dienen zunächst der Vorbereitung; zum Termin sind die erforderlichen gültigen Ausweisdokumente mitzubringen.

Bitte übersenden Sie auch die ergänzende Gesellschaftervereinbarung zur Prüfung des Zusammenhangs mit der Satzung. Sobald die Unterlagen abgestimmt sind, können wir einen Termin besprechen. Mit dieser Nachricht ist noch kein Termin bestätigt und keine Gründung beurkundet.

Mit freundlichen Grüßen
Notariatsbüro
""")

mail("08_Letzter_Stand.eml", "Frieda <frieda@topf-tacheles.example>",
     "Gründungskreis <gruendung@topf-tacheles.example>", "Unser Stand heute – bitte erst die ergänzten Entwürfe lesen", "16:40", """
Hallo zusammen,

der Anteilsplan bleibt bei 25.000 Euro und 30 / 25 / 15 / 10 / 10 / 5 / 5 Prozent. Die Zuordnung P01 bis P07 ist unverändert. Die vollständigen Personendaten stehen auf unseren sieben Datenkarten; in den beiden Verträgen sind dafür noch offene Felder vorgesehen.

Über Stimmen, besondere Mehrheiten, Lios gewünschte Sperre und die Grenzen für größere Käufe und Kredite haben wir uns heute nicht geeinigt. Auch meine Vertretungsbefugnis und die Einzelheiten zur Mitarbeit müssen wir noch besprechen. Bitte diese Lücken nicht als bereits entschiedene Regeln behandeln.

Wir warten auf die ausgefüllten und geprüften Entwürfe und lesen sie dann gemeinsam. Der Firmenname ist weiterhin nicht abschließend geprüft. Es gab keine Beurkundung, keine Kontoeröffnung und keine Einzahlung; im Handelsregister steht die GmbH noch nicht. Ein Notartermin ist noch nicht bestätigt.

Viele Grüße
Frieda
""")

print("Erzeugt: 4 EML, 1 TXT und 3 DOCX in " + str(OUT))
