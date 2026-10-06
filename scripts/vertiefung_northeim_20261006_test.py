#!/usr/bin/env python3
"""Rechen- und Paketprüfung einschließlich echter Eingabemutationen in LibreOffice."""
from pathlib import Path
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import argparse
import hashlib
import importlib.util
import io
import json
import os
import re
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from pypdf import PdfReader
from office_process import run_office
from testakte_office_pdf import office_binary
from testakte_disclaimer import NOTICE_BYTES,pdf_content_errors

ROOT=Path(__file__).resolve().parents[1]
CASE=ROOT/'testakten/bauwirtschaft-vergabeverfahren-feuerwehrhaus-northeim'
QA=Path('/tmp/bauwirtschaft-vertiefung-20261006/northeim')
OUT=QA.parent/'dist'
q=lambda n:'{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'+n

def get(n):return next(CASE.glob(str(n)+'_*.xlsx'))
def near(a,b):assert abs(float(a)-float(b))<1e-7,(a,b)
def cents(n):return Decimal(str(n)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)

def repository():
    original=json.loads((QA/'original-hashes.json').read_text())
    assert len(original)==32
    assert all(hashlib.sha256((CASE/k).read_bytes()).hexdigest()==v for k,v in original.items())
    for p in [get(43),get(44),get(45)]:
        v=load_workbook(p,data_only=True);f=load_workbook(p,data_only=False)
        assert v.properties.creator=='Klotzkette',(p.name,v.properties.creator)
        for s in f:
            assert s.print_area,(p.name,s.title)
            for row in s:
                for c in row:
                    assert c.data_type!='e',(p.name,s.title,c.coordinate,c.value)
                    if c.data_type=='f':assert v[s.title][c.coordinate].data_type!='e'
        v.close();f.close()
    w=load_workbook(get(43),data_only=True)
    for i,expected in enumerate([313600,356178,371876],6):near(w['Preisvergleich'][f'B{i}'].value,expected)
    for r in range(6,26):
        ident=w['Angebotswerte'][f'A{r}'];assert ident.data_type=='s' and re.fullmatch(r'01\.\d{3}',ident.value)
        for source,cc,outcol in [('08_angebot_leinetal.xlsx','E','D'),('10_angebot_weserklima.xlsx','F','E'),('12_angebot_harzraum.xlsx','G','F')]:
            src=load_workbook(CASE/source,data_only=True)
            assert w['Angebotswerte'][f'{cc}{r}'].value==src['LV'][f'E{r}'].value
            assert cents(w['Preisvergleich'][f'{outcol}{r+7}'].value)==cents(src['LV'][f'F{r}'].value);src.close()
    near(w['Kostenaufklärung']['B14'].value,0);w.close()
    w=load_workbook(get(44),data_only=True);near(w['Gerätewerte']['C10'].value,1.89);near(w['Gerätewerte']['D10'].value,1.68);near(w['Nachweisstand']['C14'].value,4);w.close()
    w=load_workbook(get(45),data_only=True);assert w['Termine']['B23'].value==datetime(2026,9,25);assert w['Termine']['E8'].value=='bis Termin erfasst';near(w['Unterlagen']['C15'].value,5);w.close()
    result={'originale':45,'originale_unveraendert':32,'mappen_neu':3,'blaetter_neu':7}
    print(json.dumps(result,ensure_ascii=False));return result

def mutate(source,target,sheet_number,cell,value):
    with zipfile.ZipFile(source)as src,zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED)as out:
        for item in src.infolist():
            data=src.read(item.filename)
            if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',item.filename):
                tree=ET.fromstring(data)
                if item.filename==f'xl/worksheets/sheet{sheet_number}.xml':
                    c=next(c for c in tree.iter(q('c'))if c.get('r')==cell)
                    for child in list(c):c.remove(child)
                    c.attrib.pop('t',None)
                    if value is not None:
                        if isinstance(value,str):c.set('t','inlineStr');ET.SubElement(ET.SubElement(c,q('is')),q('t')).text=value
                        else:ET.SubElement(c,q('v')).text=str(value)
                for c in tree.iter(q('c')):
                    if c.find(q('f')) is not None:
                        for child in list(c):
                            if child.tag in(q('v'),q('is')):c.remove(child)
                        c.attrib.pop('t',None)
                data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            out.writestr(item,data)

def native():
    specs=[
      (43,2,'E7',178000,'Preisvergleich','C6',3,'Preisrang ändert sich'),
      (43,2,'E7',None,'Preisvergleich','B6','Preis fehlt','Leerpreis sperrt Summe'),
      (43,2,'E7',0,'Preisvergleich','B6',235600,'Nullpreis bleibt Zahl'),
      (43,2,'E6',12800.005,'Preisvergleich','B6',313600.01,'Cent-Rundung je Position'),
      (43,2,'C9',561,'Preisvergleich','B6',313718,'Mengenänderung erreicht Summe'),
      (43,3,'B6',62000,'Kostenaufklärung','B14',1000,'Einkauf und Differenz'),
      (44,1,'D6',None,'Gerätewerte','D10','Eingabe fehlt','Fehlender Luftstrom'),
      (44,1,'D6',0,'Gerätewerte','D10','Eingabe fehlt','Nullnenner abgefangen'),
      (44,1,'D9',0,'Gerätewerte','D10',0,'Nullleistung ist numerisch'),
      (44,1,'D9',3,'Gerätewerte','D13','Grenze eingehalten','SFP exakt 1,80'),
      (44,1,'D9',3.001,'Gerätewerte','D13','Grenze überschritten','SFP knapp über 1,80'),
      (44,2,'C8','vorliegend','Nachweisstand','C16','Weitere Angaben offen','Ein Nachweis erledigt andere nicht'),
      (44,2,'C8',None,'Nachweisstand','C16','Arbeitsstand fehlt','Fehlender Nachweisstatus'),
      (45,1,'B5',46296.5,'Termine','E11','offen, Termin vorbei','Stichtag nach Frist'),
      (45,1,'C11',46295.5,'Termine','E11','bis Termin erfasst','Eingang genau zur Frist'),
      (45,1,'C11',46295.5000115741,'Termine','E11','nach Termin erfasst','Eingang eine Sekunde später'),
      (45,2,'D7','erledigt','Unterlagen','C17','Arbeiten offen','Ein Arbeitsschritt erledigt andere nicht'),
      (45,2,'D7',None,'Unterlagen','C17','Arbeitsstand fehlt','Leerer Arbeitsstand'),
    ]
    before={p.name:hashlib.sha256(p.read_bytes()).hexdigest()for p in [get(43),get(44),get(45)]}
    results=[]
    with tempfile.TemporaryDirectory(prefix='northeim-mutation-')as td:
        temp=Path(td);out=temp/'out';out.mkdir()
        paths=[]
        for i,(book,sheet,cell,value,*_)in enumerate(specs):
            dest=temp/f'test-{i:02}.xlsx';mutate(get(book),dest,sheet,cell,value);paths.append(dest)
        for start in range(0,len(paths),6):
            cmd=[office_binary(),f'-env:UserInstallation={(temp/"profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),*[str(p)for p in paths[start:start+6]]]
            code,log=run_office(cmd,timeout=180,env=os.environ.copy());assert code==0,log
        for i,(_,_,_,_,sheet,cell,expected,meaning)in enumerate(specs):
            w=load_workbook(out/paths[i].name,data_only=True);actual=w[sheet][cell].value
            if isinstance(expected,str):assert actual==expected,(meaning,actual,expected)
            else:near(actual,expected)
            results.append({'pruefung':meaning,'erwartet':expected,'erhalten':actual});w.close()
    assert before=={p.name:hashlib.sha256(p.read_bytes()).hexdigest()for p in [get(43),get(44),get(45)]}
    (QA/'native-mutations.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
    print(f'{len(results)} native Mutationstests bestanden; Originalmappen unverändert.');return results

def packages():
    manifests=json.loads((QA/'manifest.json').read_text());assert len(manifests)==45
    total=CASE/'gesamt-pdf'/f'{CASE.name}_gesamt.pdf';assert len(PdfReader(total).pages)==sum(x['pages']for x in manifests)
    assert not pdf_content_errors(total.read_bytes())
    for suffix,count in [('',47),('-einzelpdfs',46)]:
        file=OUT/f'testakte-{CASE.name}{suffix}.zip'
        with zipfile.ZipFile(file)as z:
            assert len(z.namelist())==count;assert z.namelist()[0]=='README.txt';assert z.read('README.txt').startswith(NOTICE_BYTES)
            assert not any('/'in n or n.endswith(('.md','.yaml','.json','.ndjson','.py','.mjs'))for n in z.namelist())
            for n in z.namelist():
                if n.endswith('.pdf'):assert not pdf_content_errors(z.read(n)),n
            if not suffix:
                for p in CASE.iterdir():
                    if p.name[:2].isdigit():assert z.read(p.name)==p.read_bytes()
    print(f'45 Einzel-PDFs und Gesamt-PDF mit {len(PdfReader(total).pages)} Seiten sowie beide Archive geprüft.')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--native',action='store_true');p.add_argument('--packages',action='store_true');a=p.parse_args();repository()
    if a.native:native()
    if a.packages:packages()
