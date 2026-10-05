#!/usr/bin/env python3
"""Begrenzte Integritätsprüfungen der neuen Geburtsschadenakte."""
import unittest,sys
from pathlib import Path
from email import policy
from email.parser import BytesParser
from openpyxl import load_workbook
from docx import Document
from geburtsschaden_akte_daten import SLUG,DOCS,ATTACHMENTS,XLSX,CORE
ROOT=Path(__file__).resolve().parents[1];CASE=ROOT/'testakten'/SLUG
class GeburtsschadenTests(unittest.TestCase):
 def test_originalbestand(self):
  names={d['file'] for d in DOCS}|{XLSX};actual={p.name for p in CASE.iterdir() if p.suffix.lower() in {'.eml','.docx','.xlsx','.pdf','.txt'}}
  self.assertEqual(actual,names);self.assertEqual(len(names),18);self.assertLessEqual(len(CORE),6)
 def test_echte_anlagen(self):
  for name,attachments in ATTACHMENTS.items():
   msg=BytesParser(policy=policy.default).parsebytes((CASE/name).read_bytes());parts={p.get_filename():p.get_payload(decode=True) for p in msg.iter_attachments()}
   self.assertEqual(set(parts),set(attachments))
   for filename in attachments:self.assertEqual(parts[filename],(CASE/filename).read_bytes())
 def test_aktive_formeln(self):
  w=load_workbook(CASE/XLSX,data_only=False);self.assertEqual(w['Ausgleichsmodell']['B14'].value,'=MIN(MAX(0,B11-B8),B9)')
  self.assertEqual(w['Ausgleichsmodell']['B13'].value,'=MAX(0,B11-B12)');self.assertEqual(w['Vergleich']['B15'].value,'=B13-C13');self.assertEqual(w['Zahlungen']['C13'].value,'=SUM(C7:C11)')
 def test_gecachte_ergebnisse(self):
  w=load_workbook(CASE/XLSX,data_only=True)
  for sheet,cell,expected in [('Vergleich','B13',13200000),('Vergleich','C13',1200000),('Vergleich','B15',12000000),('Zahlungen','C19',4700000),('Ausgleichsmodell','B13',11700000),('Ausgleichsmodell','B14',3200000),('Ausgleichsmodell','B15',8500000),('Ausgleichsmodell','B23',650000)]:self.assertAlmostEqual(w[sheet][cell].value,expected,places=2)
 def test_rollen_und_vorbehalt(self):
  text='\n'.join(p.text for p in Document(CASE/'11_Teilvergleich.docx').paragraphs)
  for token in ['ausschließlich für Nora','nicht vorhersehbar','Sozialleistungsträger','15. Juni 2048']:self.assertIn(token,text)
  self.assertNotIn('genehmigt durch',text)
 def test_modell_ausdruecklich_begrenzt(self):
  text='\n'.join(p.text for p in Document(CASE/'14_Interne_Modellvereinbarung.docx').paragraphs)
  for token in ['keine reale Mitgliedschaft','ursprünglichen anrechenbaren Gesamtschaden','kein Selbstbehalt der Klinik']:self.assertIn(token,text)
if __name__=='__main__':unittest.main()
