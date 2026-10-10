#!/usr/bin/env python3
"""Bestands-, Workflow- und Exportprüfungen des VVT-Komponentenpakets."""
from pathlib import Path
from collections import Counter
import argparse
import hashlib
import importlib.util
import io
import json
import re
import shutil
import sys
import tempfile
import unittest
import zipfile
from email import policy
from email.parser import BytesParser
from openpyxl import load_workbook
from pypdf import PdfReader
import pdfplumber
import yaml
from testakte_einzelpdf_common import document_arcname_pairs
from testakte_zip_common import working_dump_archive_pairs
from testakte_disclaimer import NOTICE_BYTES, NOTICE_DE, NOTICE_EN
from release_routing import case_asset_url, plugin_asset_url
from vvt_workbook_pdf import render as render_workbook

ROOT=Path(__file__).resolve().parents[1]
NAME='verarbeitungsverzeichnis'; PLUGIN=ROOT/NAME; DIST=None
CASES=('vvt-handwerk-finkenbeil-erfurt','vvt-cloudservice-wolkengarn-berlin','vvt-praxis-rosenquell-bamberg')

def module(path):
    spec=importlib.util.spec_from_file_location(path.stem.replace('-','_'),path)
    out=importlib.util.module_from_spec(spec); spec.loader.exec_module(out); return out

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

class VVTComponentTests(unittest.TestCase):
    def test_manifests_skills_commands_and_local_links(self):
        for rel in ('plugin.json','.claude-plugin/plugin.json','.codex-plugin/plugin.json'):
            obj=json.loads((PLUGIN/rel).read_text()); self.assertEqual((obj['name'],obj['version']),(NAME,'1.0.0'))
        skills=list((PLUGIN/'skills').glob('*/SKILL.md')); self.assertEqual(len(skills),8)
        for path in skills:
            content=path.read_text(); meta=yaml.safe_load(content.split('---',2)[1])
            self.assertEqual(set(meta),{'name','description'}); self.assertEqual(meta['name'],path.parent.name)
            self.assertLessEqual(len(meta['description']),360); self.assertNotIn('§',meta['description'])
            for i in range(1,7): self.assertRegex(content,rf'(?m)^#{{1,3}} {i}[. ]')
            for required in ('Times New Roman','11 pt','ausformuliert'):self.assertIn(required,content)
        commands=list((PLUGIN/'commands').glob('*.md'))
        self.assertEqual({p.stem for p in commands},{'vvt','vvt-aenderung','vvt-risiko','vvt-export'})
        for path in commands:
            parts=path.read_text().split('---',2); self.assertEqual(parts[0],'')
            self.assertTrue(yaml.safe_load(parts[1])['description']);self.assertIn('$ARGUMENTS',parts[2])
        for path in PLUGIN.rglob('*.md'):
            for target in re.findall(r'\]\(([^ )]+)\)',path.read_text()):
                if target.startswith(('http:','https:','#')):continue
                dest=(path.parent/target.split('#')[0]).resolve()
                self.assertTrue(dest.is_relative_to(PLUGIN.resolve()),f'{path}: {target}')
                self.assertTrue(dest.exists(),f'{path}: {target}')

    def test_prompt_pairs_and_frozen_review_hashes(self):
        profile=json.loads((ROOT/'quality/evals/verarbeitungsverzeichnis.json').read_text())
        for kind,key in [('werkstatt','workshop_review'),('schnellstart','mini_review'),('hauptproblem','focus_review')]:
            path=PLUGIN/f'{NAME}-{kind}.md'
            self.assertEqual(path.read_bytes(),path.with_suffix('.txt').read_bytes())
            if kind!='werkstatt':self.assertLessEqual(path.stat().st_size,7500)
            self.assertEqual(profile[key]['prompt_sha256' if kind=='hauptproblem' else 'sha256'],sha(path))
        self.assertEqual(profile['focus_review']['skill_sha256'],sha(PLUGIN/'skills/verarbeitungsverzeichnis-steuern/SKILL.md'))

    def test_case_inputs_registers_and_original_workbook_roundtrips(self):
        vvt=module(PLUGIN/'scripts/vvt.py'); roles=set()
        for case in CASES:
            folder=ROOT/'testakten'/case
            pairs=document_arcname_pairs(folder)
            self.assertEqual(len(pairs),14,case)
            before=vvt.read_register(folder/'02_Registerbestand.json'); after=vvt.read_register(folder/'03_Register_Aenderungsstand.json')
            self.assertEqual(before['register_id'],after['register_id'])
            self.assertGreater(after['revision'],before['revision'])
            self.assertTrue({a['id'] for a in before['activities']} <= {a['id'] for a in after['activities']})
            self.assertNotEqual(vvt.digest(before),vvt.digest(after))
            roles.update(a['role'] for a in after['activities'])
            self.assertTrue(any(a['screening']['high_risk']=='offen' for a in after['activities']))
            with tempfile.TemporaryDirectory() as tmp:
                local=Path(tmp)/'register.json';shutil.copyfile(folder/'03_Register_Aenderungsstand.json',local)
                result=vvt.import_xlsx(local,folder/'04_Verzeichnis.xlsx','Clara Winde','Unveränderten Export abgleichen')
                self.assertEqual(result,after)
            book=load_workbook(folder/'04_Verzeichnis.xlsx',data_only=False)
            self.assertEqual(set(book.sheetnames),{*vvt.SHEETS,'Uebersicht','Organisation','Entscheidungen','_meta'})
            for tab in book:
                for row in tab:
                    for cell in row:self.assertNotEqual(cell.data_type,'e')
        self.assertEqual(roles,{'controller','processor'})

    def test_correspondence_and_case_download_notices(self):
        for case in CASES:
            folder=ROOT/'testakten'/case
            emails=list(folder.glob('*.eml'));self.assertEqual(len(emails),4)
            for path in emails:
                message=BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                self.assertFalse(message.defects); self.assertTrue(message['Subject']);self.assertTrue(message['Date'])
                self.assertTrue(message['From']);self.assertTrue(message['To']);self.assertTrue(message['Message-ID'])
                self.assertGreater(len(message.get_body(preferencelist=('plain',)).get_content().split()),40)
            readme=(folder/'README.md').read_text()
            for notice in (NOTICE_DE,NOTICE_EN):
                for line in notice.splitlines():
                    if line.strip():self.assertIn(line.strip(),readme)
            self.assertIn('verarbeitungsverzeichnis-v1.0.0',readme)
            for suffix in ('.zip','-einzelpdfs.zip','_gesamt.pdf'):self.assertIn(case+suffix,readme)

    def test_neutral_starter_exports_validate_and_share_one_register(self):
        vvt=module(PLUGIN/'scripts/vvt.py'); folder=PLUGIN/'templates/startregister'
        self.assertTrue(folder.is_dir())
        files=list(folder.glob('*.json'));self.assertEqual(len(files),1)
        document=vvt.read_register(files[0])
        for suffix in ('xlsx','docx','xml','html'):self.assertEqual(len(list(folder.glob('*.'+suffix))),1)
        meta,parsed=vvt.parse_xml(next(folder.glob('*.xml')));self.assertEqual(parsed,document)
        with tempfile.TemporaryDirectory() as tmp:
            dest=Path(tmp)/'register.json';shutil.copyfile(files[0],dest)
            self.assertEqual(vvt.import_xlsx(dest,next(folder.glob('*.xlsx')),'Clara Winde','Vorlage prüfen'),document)

    def test_component_routes(self):
        self.assertIn('/verarbeitungsverzeichnis-v1.0.0/',plugin_asset_url(NAME))
        for case in CASES:
            for suffix in ('','-einzelpdfs'):self.assertIn('/verarbeitungsverzeichnis-v1.0.0/',case_asset_url(case,suffix))

    def test_workbook_reading_view_preserves_cells_and_legibility(self):
        for case in CASES:
            path=ROOT/'testakten'/case/'04_Verzeichnis.xlsx';data=render_workbook(path)
            text='\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(data)).pages)
            compact=re.sub(r'\s+','',text)
            for sheet in load_workbook(path,data_only=False):
                self.assertIn(sheet.title,text)
                for row in sheet:
                    for cell in row:
                        if cell.value is not None:
                            value=cell.value.isoformat() if hasattr(cell.value,'isoformat') else str(cell.value)
                            self.assertIn(re.sub(r'\s+','',value),compact,f'{case} {sheet.title}!{cell.coordinate}')
            with pdfplumber.open(io.BytesIO(data)) as doc:
                for page in doc.pages:
                    for char in page.chars:
                        self.assertGreaterEqual(char['size'],8)
                        self.assertGreaterEqual(char['x0'],45)
                        self.assertLessEqual(char['x1'],page.width-45)
                        self.assertLessEqual(char['bottom'],page.height-15)

    def require_dist(self):
        if DIST is None:self.skipTest('Release-Dateien erst nach --dist prüfen.')

    def test_dist_exact_inventory_checksums_and_plugin_bytes(self):
        self.require_dist(); builder=module(ROOT/'scripts/build-verarbeitungsverzeichnis-release.py')
        expected=builder.asset_names()|{'checksums-sha256.txt'}
        self.assertEqual({p.name for p in DIST.iterdir()},expected);self.assertEqual(len(expected),22)
        for line in (DIST/'checksums-sha256.txt').read_text().splitlines():
            digest,name=line.split('  ',1);self.assertEqual(digest,sha(DIST/name))
        for suffix,prefix in [('', ''),('-portable',NAME+'/')]:
            with zipfile.ZipFile(DIST/f'{NAME}{suffix}.zip') as archive:
                sources=builder.plugin_sources(); expected_names={prefix+p.relative_to(PLUGIN).as_posix() for p in sources}|{prefix+'LICENSE'}
                self.assertEqual(set(archive.namelist()),expected_names)
                for path in sources:self.assertEqual(archive.read(prefix+path.relative_to(PLUGIN).as_posix()),path.read_bytes())
                self.assertFalse(any('/testakten/' in n or '/templates/' in n for n in archive.namelist()))
        with zipfile.ZipFile(DIST/f'{NAME}-vorlagen.zip') as archive:
            for path in (PLUGIN/'templates').rglob('*'):
                if path.is_file():self.assertEqual(archive.read(path.relative_to(PLUGIN/'templates').as_posix()),path.read_bytes())

    def test_case_zips_original_bytes_and_pdf_page_text_identity(self):
        self.require_dist()
        for case in CASES:
            folder=ROOT/'testakten'/case
            with zipfile.ZipFile(DIST/f'testakte-{case}.zip') as archive:
                self.assertTrue(archive.read('README.txt').startswith(NOTICE_BYTES))
                for path,name in working_dump_archive_pairs(folder, include_gesamt_pdf=True):self.assertEqual(archive.read(name),path.read_bytes())
            combined=PdfReader(DIST/f'{case}_gesamt.pdf');pages=[]
            with zipfile.ZipFile(DIST/f'testakte-{case}-einzelpdfs.zip') as archive:
                self.assertTrue(archive.read('README.txt').startswith(NOTICE_BYTES))
                self.assertEqual(set(archive.namelist()),{n for _,n in document_arcname_pairs(folder)}|{'README.txt'})
                for _,name in document_arcname_pairs(folder):pages.extend(p.extract_text() or '' for p in PdfReader(io.BytesIO(archive.read(name))).pages)
            self.assertEqual(pages,[p.extract_text() or '' for p in combined.pages])
            for text in pages:
                self.assertNotIn('Diese Testakte wurde mit KI generiert',text)
                self.assertNotIn('This test case file was generated with AI',text)
                self.assertNotIn('\ufffd',text)

    def test_prompt_pdf_text_complete(self):
        self.require_dist()
        for kind in ('werkstatt','schnellstart','hauptproblem'):
            stem=f'{NAME}-{kind}';source=(PLUGIN/f'{stem}.md').read_text()
            for suffix in ('md','txt'):self.assertEqual((DIST/f'{stem}.{suffix}').read_bytes(),(PLUGIN/f'{stem}.md').read_bytes())
            text='\n'.join(p.extract_text() or '' for p in PdfReader(DIST/f'{stem}.pdf').pages)
            words=lambda value:Counter(re.findall(r'[\w§]+',value.casefold()))
            self.assertFalse(words(source)-words(text),kind)
            self.assertNotIn('\ufffd',text);self.assertNotIn('\u25a0',text)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--dist',type=Path)
    args,extra=parser.parse_known_args();DIST=args.dist
    unittest.main(argv=[sys.argv[0],*extra])
