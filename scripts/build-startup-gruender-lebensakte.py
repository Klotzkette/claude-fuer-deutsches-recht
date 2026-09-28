#!/usr/bin/env python3
"""Baut native Korrespondenz und Unterlagen der Schnittflug-Gründungsakte.

Datenquelle: scripts/data/startup-gruender/kommunikation.json.
DOCX verwenden Times New Roman 11 pt; PNG bilden fiktive Teamoberflächen
in Bildschirmtypografie ab. Keine Fremdkonten, amtlichen Vollzüge oder Fotos.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
from email import policy
from email.message import EmailMessage
from email.utils import format_datetime, getaddresses
import hashlib
import io
import json
import mimetypes
from pathlib import Path
import zipfile

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'scripts/data/startup-gruender/kommunikation.json'
FACTS = ROOT / 'scripts/data/startup-gruender/fallstamm.json'
FIXED_DATE = datetime(2026, 9, 28, 9, 0, tzinfo=timezone.utc)


def save_docx(row: dict, folder: Path) -> None:
    doc = Document()
    layout = row.get('layout', {})
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(layout.get('vertical_margin_cm', 2.0))
    sec.left_margin = sec.right_margin = Cm(2.25)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = layout.get('line_spacing', 1.08)
    normal.paragraph_format.space_after = Pt(layout.get('paragraph_space_after_pt', 6))
    normal.paragraph_format.widow_control = True
    for style_name, size in [('Title', 18), ('Subtitle', 11), ('Heading 1', 13), ('Heading 2', 11)]:
        style = doc.styles[style_name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10 if style_name.startswith('Heading') else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        if style_name.startswith('Heading'):
            style.font.bold = True
    for node in list(doc.styles.element.iter()):
        if node.tag == qn('w:pBdr'):
            node.getparent().remove(node)
        elif node.tag == qn('w:rFonts'):
            for key in list(node.attrib):
                if 'theme' in key.lower():
                    del node.attrib[key]
            node.set(qn('w:ascii'), 'Times New Roman')
            node.set(qn('w:hAnsi'), 'Times New Roman')
    doc.core_properties.author = row['author']
    doc.core_properties.title = row['title']
    doc.core_properties.subject = 'Schnittflug Robotics – Gründungsunterlagen'
    doc.core_properties.created = doc.core_properties.modified = FIXED_DATE
    doc.add_paragraph(row['title'], 'Title')
    p = doc.add_paragraph(row['author'], 'Subtitle')
    doc.add_paragraph(row['date'])
    doc.add_paragraph(row['intro'])
    for section in row['sections']:
        heading = doc.add_paragraph(section['title'], 'Heading 1')
        if section.get('page_break_before'):
            heading.paragraph_format.page_break_before = True
        for paragraph in section['paragraphs']:
            doc.add_paragraph(paragraph)
    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    footer.alignment = 2
    r = footer.add_run('Schnittflug · ')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    stream = io.BytesIO()
    doc.save(stream)
    # Package timestamps are fixed, so repeated builds preserve attachment hashes.
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(folder / row['file'], 'w', compression=zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 28, 9, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))


def ui_font(size: int, bold: bool = False):
    file = 'Arial Bold.ttf' if bold else 'Arial.ttf'
    candidates = [Path('/System/Library/Fonts/Supplemental') / file,
                  Path('/usr/share/fonts/truetype/dejavu') / ('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf')]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    raise RuntimeError('Für lesbare UI-Bilder wird Arial oder DejaVu Sans benötigt.')


def wrapped(draw, text, font, width):
    result = []
    for line in text.split('\n'):
        current = ''
        for word in line.split():
            candidate = (current + ' ' + word).strip()
            if current and draw.textlength(candidate, font=font) > width:
                result.append(current)
                current = word
            else:
                current = candidate
        result.append(current)
    return result


def screenshot(row: dict, folder: Path):
    im = Image.new('RGB', (1400, 1080), '#eef2f6')
    draw = ImageDraw.Draw(im)
    fonts = {s: ui_font(s) for s in (21, 24, 27)}
    bold = {s: ui_font(s, True) for s in (22, 27, 36)}
    draw.rectangle((0, 0, 1400, 142), fill='#142f43')
    draw.text((48, 32), row['title'], font=bold[36], fill='white')
    draw.text((48, 91), row['subtitle'], font=fonts[24], fill='#d5e4ee')
    if 'cards' in row:
        for i, (status, title, body) in enumerate(row['cards']):
            x = 42 + (i % 2) * 676
            y = 178 + (i // 2) * 400
            draw.rounded_rectangle((x, y, x + 640, y + 364), radius=14, fill='white', outline='#ccd6de', width=2)
            draw.text((x + 24, y + 24), status, font=bold[22], fill='#276379')
            draw.text((x + 24, y + 69), title, font=bold[27], fill='#142f43')
            yy = y + 125
            for line in wrapped(draw, body, fonts[27], 592):
                if yy + 33 > y + 345:
                    raise ValueError(f'UI-Karte überläuft: {row["file"]} {title}')
                draw.text((x + 24, yy), line, font=fonts[27], fill='#223b4d')
                yy += 38
    else:
        y = 174
        for time, name, body in row['messages']:
            lines = wrapped(draw, body, fonts[27], 1158)
            h = 62 + len(lines) * 35
            draw.rounded_rectangle((44, y, 1356, y + h), radius=12, fill='white', outline='#cdd8e0', width=1)
            draw.text((69, y + 15), name, font=bold[22], fill='#276379')
            draw.text((1242, y + 16), time, font=fonts[21], fill='#657b8a')
            for index, line in enumerate(lines):
                draw.text((69, y + 53 + index * 35), line, font=fonts[27], fill='#223b4d')
            y += h + 15
        if y > 1006:
            raise ValueError(f'Chat-Bild überläuft: {row["file"]}')
    draw.text((46, 1034), 'Schnittflug · Teamablage · Stand wie angezeigt', font=fonts[21], fill='#546c7e')
    im.save(folder / row['file'])


def save_mail(row, index, records, people, folder):
    def address(value):
        return people.get(value, value)
    msg = EmailMessage(policy=policy.SMTP)
    for key in ('sender', 'to'):
        value = address(row[key])
        for _, addr in getaddresses([value]):
            if not addr.endswith('.example'):
                raise ValueError(f'Nicht reservierte Kontaktadresse: {addr}')
    msg['From'] = address(row['sender'])
    msg['To'] = address(row['to'])
    if row.get('cc'):
        msg['Cc'] = ', '.join(address(x) for x in row['cc'])
    msg['Date'] = format_datetime(datetime.fromisoformat(row['date']))
    msg['Subject'] = row['subject']
    msg['Message-ID'] = f'<schnittflug-2026-{index:02d}@schnittflug.example>'
    if row.get('reply'):
        reply = row['reply']
        msg['In-Reply-To'] = f'<schnittflug-2026-{reply:02d}@schnittflug.example>'
        ancestors = [reply]
        while records[ancestors[-1] - 1].get('reply'):
            ancestors.append(records[ancestors[-1] - 1]['reply'])
        msg['References'] = ' '.join(f'<schnittflug-2026-{n:02d}@schnittflug.example>' for n in reversed(ancestors))
    msg['X-Mailer'] = 'Mail 16.0' if index % 3 else 'Büropost Desktop 4.2'
    body = row['body']
    if row.get('reply'):
        prior = records[row['reply'] - 1]
        quoted = prior['body'].split('\n\n')[1]
        body += f'\n\nAm {datetime.fromisoformat(prior["date"]).strftime("%d.%m.%Y um %H:%M")} schrieb {address(prior["sender"])}:\n> ' + quoted.replace('\n', '\n> ') + '\n'
    msg.set_content(body, charset='utf-8', cte='quoted-printable')
    for filename in row['attachments']:
        path = folder / filename
        if path.parent != folder or not path.is_file():
            raise ValueError(f'Unzulässiger oder fehlender Anhang: {filename}')
        mime = mimetypes.guess_type(filename)[0] or 'application/octet-stream'
        maintype, subtype = mime.split('/', 1)
        msg.add_attachment(path.read_bytes(), maintype=maintype, subtype=subtype, filename=filename)
    if row['attachments']:
        msg.set_boundary(f'=_schnittflug_2026_{index:02d}')
    (folder / row['file']).write_bytes(msg.as_bytes())


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path)
    args = parser.parse_args()
    data = json.loads(DATA.read_text(encoding='utf-8'))
    facts = json.loads(FACTS.read_text(encoding='utf-8'))
    assert data['case_slug'] == facts['case_slug']
    assert len(facts['founders']) == 7
    assert sum(p['nominal_eur'] for p in facts['founders']) == 25000
    folder = args.output_dir or ROOT / 'testakten' / data['case_slug']
    folder.mkdir(parents=True, exist_ok=True)
    records = [*data['documents'], *data['mails'], *data['chats'], *data['screenshots']]
    names = [row['file'] for row in records]
    if len(names) != len(set(names)) or any(Path(n).name != n for n in names):
        raise ValueError('Aktennamen sind nicht eindeutig und flach.')
    for row in data['documents']:
        save_docx(row, folder)
    for row in data['screenshots']:
        screenshot(row, folder)
    for row in data['chats']:
        text = row['title'] + '\n' + row['date'] + '\n\n'
        text += '\n\n'.join(f'[{date}] {name}: {body}' for date, name, body in row['messages']) + '\n'
        (folder / row['file']).write_text(text, encoding='utf-8')
    for index, row in enumerate(data['mails'], 1):
        save_mail(row, index, data['mails'], data['people'], folder)
    summary = {'files': len(names), 'docx': len(data['documents']), 'eml': len(data['mails']), 'mime_attachments': sum(len(r['attachments']) for r in data['mails']), 'chat_txt': len(data['chats']), 'screenshots_png': len(data['screenshots'])}
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == '__main__':
    main()
