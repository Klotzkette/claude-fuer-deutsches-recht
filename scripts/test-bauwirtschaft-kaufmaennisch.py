#!/usr/bin/env python3
"""Belegintegrität, Exportabgleich und echte Office-Neuberechnung. Autor: Klotzkette."""
import argparse
import contextlib
import csv
import hashlib
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal
from datetime import datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from pypdf import PdfReader
from testakte_office_pdf import office_binary
from testakte_disclaimer import notice_text_errors

ROOT = Path(__file__).resolve().parents[1]
ASSETS = None
QA = None
CASES = ['bauwirtschaft-baumanagement-werkhalle-warendorf','bauwirtschaft-buchhaltung-bauunternehmen-bad-salzuflen']
NOTICE = ('Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\n'
          'This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.')
NS = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RESULTS = {'author':'Klotzkette','checks':[], 'office_mutations':[], 'live_model_tests':False}


def check(value, label):
    assert value, label
    RESULTS['checks'].append(label)


def amount(value):
    return Decimal(value.replace('.','').replace(',','.'))


def csvrows(case, name):
    with (ROOT/'testakten'/case/name).open(encoding='utf-8-sig',newline='') as stream:
        return list(csv.DictReader(stream,delimiter=';'))


def near(a,b):
    if isinstance(a,datetime):a=(a-datetime(1899,12,30)).total_seconds()/86400
    if isinstance(b,datetime):b=(b-datetime(1899,12,30)).total_seconds()/86400
    return abs(float(a)-float(b)) < .005


def pdftext(source):
    return '\n'.join(p.extract_text() or '' for p in PdfReader(source).pages)


def normalized(text):
    return ''.join(c for c in unicodedata.normalize('NFKC',str(text)).casefold() if c.isalnum())


def source_fragments(source):
    if source.suffix=='.pdf':
        return [page.extract_text() or '' for page in PdfReader(source).pages]
    if source.suffix=='.docx':
        doc=Document(source)
        return [p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]
    if source.suffix=='.csv':
        with source.open(encoding='utf-8-sig',newline='') as stream:
            return [c for row in csv.reader(stream,delimiter=';') for c in row]
    if source.suffix=='.txt':
        return source.read_text(encoding='utf-8').splitlines()
    if source.suffix=='.eml':
        message=BytesParser(policy=policy.default).parsebytes(source.read_bytes())
        return [str(message.get(key,'')) for key in ('From','To','Cc','Subject')]+message.get_body(preferencelist=('plain',)).get_content().splitlines()
    if source.suffix=='.xlsx':
        workbook=load_workbook(source,data_only=True)
        # Anzeigeformate für Zahlen und Datumswerte unterscheiden sich je Office-Version.
        fragments=[c.value for sheet in workbook for row in sheet for c in row if isinstance(c.value,str)]
        workbook.close()
        return fragments
    return []


def image_digest(image):
    rgb=image.convert('RGB')
    return rgb.size,hashlib.sha256(rgb.tobytes()).hexdigest()


def exported_content(reader,source,label):
    text='\n'.join(page.extract_text() or '' for page in reader.pages)
    content=normalized(text)
    missing=[part for part in source_fragments(source) if normalized(part) and normalized(part) not in content]
    check(not missing,label+' vollständiger Beleginhalt'+(': '+repr(missing[:2]) if missing else ''))
    if source.suffix=='.png':
        expected=image_digest(Image.open(source))
        check(any(image_digest(image.image)==expected for page in reader.pages for image in page.images),label+' vollständiges Originalbild')
    if source.name in ('19_Bankumsatz.csv','25_Projektkosten_und_Bank.xlsx'):
        for value in ('60.000,00','1.785,00'):
            check(normalized(value) in content,label+' Bankbetrag '+value)
        if source.suffix=='.xlsx':
            check('8500000' in content,label+' Bankanfangsbestand 85.000,00 EUR')
            check('7921180' in content,label+' Bankendsaldo 79.211,80 EUR')
    if source.name=='55_Ergaenzungsabgleich.xlsx':
        for value in ('9.999,00','19.580,00','1.520,00','38.218,80','11.781,20'):
            check(normalized(value) in content,label+' Ergänzungsbetrag '+value)


def exported_pdf(reader,label):
    check(bool(reader.pages),label+' nicht leeres PDF')
    text='\n'.join(page.extract_text() or '' for page in reader.pages)
    check(NOTICE.splitlines()[0] not in text and 'This test case file' not in text,label+' kein Warntext im PDF')
    check(all(abs(sorted((float(p.mediabox.width),float(p.mediabox.height)))[0]-595.28)<4 and
              abs(sorted((float(p.mediabox.width),float(p.mediabox.height)))[1]-841.89)<4
              for p in reader.pages),label+' A4-Seiten')


def integrity():
    war,bad=CASES
    journal=csvrows(war,'11_Belegjournal.csv')
    check(sum(amount(r['Netto_EUR']) for r in journal)==1200000,'Warendorf Ist 1.200.000 EUR')
    check(sum(amount(r['USt_EUR']) for r in journal)==228000,'Warendorf Steueranteil 228.000 EUR')
    check(len({r['Beleg'] for r in journal})==14,'Warendorf 14 eindeutige Kostenbelege')
    orders=csvrows(war,'12_Bestellbuch.csv')
    check(sum(amount(r['Bestellt_netto_EUR']) for r in orders)==2073000,'Warendorf Bestellung 2.073.000 EUR')
    pay=csvrows(war,'21_Zahlungsplan.csv')
    check(sum(amount(r['Zahlung_brutto_EUR']) for r in pay)==1387540,'Warendorf Restzahlungen 1.387.540 EUR brutto')
    check(2073000-1200000==873000,'Warendorf Obligo ohne Ist-Doppelzählung 873.000 EUR')
    inv=csvrows(bad,'22_Eingangsjournal.csv')
    check(len(inv)==5 and len({r['Buchung'] for r in inv})==5,'Bad Salzuflen fünf eindeutige Buchungen')
    check(sum(r['Lieferantenbeleg']=='LB-260908' for r in inv)==1,'Rechnungskopie nur einmal gebucht')
    check(sum(amount(r['Netto_EUR']) for r in inv)==39300,'Bad Salzuflen Belegnetto 39.300 EUR')
    check(sum(amount(r['Belegbetrag_EUR']) for r in inv)==42207,'Bad Salzuflen Belegbeträge 42.207 EUR')
    bank=csvrows(bad,'19_Bankumsatz.csv')
    balance=Decimal(85000)+sum(amount(r['Haben_EUR'])-amount(r['Soll_EUR']) for r in bank)
    check(balance==Decimal('79211.80') and len(bank)==8,'Acht Bankbuchungen ergeben 79.211,80 EUR')
    check(next(r for r in bank if r['Bankreferenz']=='BK-0924C')['Verwendungszweck']=='Mietpark 09','Unzugeordnete Überweisung 1.785 EUR bleibt separat')
    opos=csvrows(bad,'23_OPOS_25_09.csv')
    for row in opos:
        check(amount(row['Belegbetrag_EUR'])-amount(row['Zahlung_EUR'])-amount(row['Skonto_EUR'])-amount(row['Gutschrift_EUR'])==amount(row['Offen_EUR']), 'OPOS-Gleichung '+row['Beleg'])
    check(sum(amount(r['Offen_EUR']) for r in opos)==3342,'OPOS 3.342 EUR einschließlich Sicherheit')
    check(Decimal(13090)*Decimal('.02')==Decimal('261.80'),'Skonto auf Betrag nach Gutschrift')
    hours=csvrows(bad,'18_Projektstunden_August.csv')
    check(sum(amount(r['Stunden'])*amount(r['Stundenlohn_EUR'])+amount(r['AG_Belastung_EUR']) for r in hours)==25400,'Augustlohn mit AG-Belastung 25.400 EUR')
    check(39300-220+25400==64480,'Projektkosten nach Nettoskonto 64.480 EUR')
    d=ROOT/'testakten'/bad
    original=pdftext(d/'08_Rechnung_LB_260908.pdf').strip()
    copy=pdftext(d/'09_Lieferantenmail_Anhang_908.pdf')
    check(original in copy and 'Posteingang 16.09.2026 10:24 Uhr' in copy,'Rechnungskern unverändert mit Eingangsstempel')
    check((d/'08_Rechnung_LB_260908.pdf').read_bytes()!=(d/'09_Lieferantenmail_Anhang_908.pdf').read_bytes(),'Eingangskopie eigenständige Datei')
    check('48b' in pdftext(d/'06_Freistellungsbescheinigung.pdf') and '13b' not in pdftext(d/'06_Freistellungsbescheinigung.pdf'),'Bauabzugsfreistellung eigenständig')
    check('USt 1 TG' in pdftext(d/'07_Bescheinigung_Bauleistungen.pdf') and '48b' not in pdftext(d/'07_Bescheinigung_Bauleistungen.pdf'),'USt 1 TG eigenständig')


def sources_and_exports():
    for case in CASES:
        directory=ROOT/'testakten'/case
        sources=sorted(p for p in directory.iterdir() if p.is_file() and p.name[:2].isdigit())
        expected_count=57 if case==CASES[1] else 30
        check(len(sources)==expected_count,case+f' {expected_count} Originale')
        check({p.suffix for p in sources}=={'.pdf','.docx','.xlsx','.csv','.txt','.eml','.png'},case+' vollständiger Formatmix')
        check(len(list(directory.glob('*.xlsx')))==(3 if case==CASES[1] else 2) and len(list(directory.glob('*.png')))==2,case+' Mappen und zwei PNGs vollständig')
        for p in sources:
            if p.suffix=='.png':
                im=Image.open(p);check(im.format=='PNG' and min(im.size)>=900,p.name+' echte hochauflösende PNG-Datei')
            elif p.suffix=='.csv':
                with p.open(encoding='utf-8-sig',newline='') as h: rows=list(csv.reader(h,delimiter=';'))
                check(all(len(r)==len(rows[0]) for r in rows) and len(rows)>=3,p.name+' vollständige CSV-Zeilen')
            elif p.suffix=='.docx':
                d=Document(p)
                check(d.core_properties.author=='Klotzkette',p.name+' Autor')
                check(d.core_properties.language=='de-DE',p.name+' Sprache')
                for i,par in enumerate(d.paragraphs):
                    if par.style.style_id in ('Heading1','Heading2'):
                        check(not d.paragraphs[i+1].text.strip(),p.name+' Überschrift mit Leerabsatz')
            elif p.suffix=='.pdf':
                check(PdfReader(p).metadata.author=='Klotzkette',p.name+' Original-PDF-Autor')
            elif p.suffix=='.eml':
                em=BytesParser(policy=policy.default).parsebytes(p.read_bytes())
                check(all(em.get(k) for k in ['From','To','Date','Subject','Message-ID','MIME-Version','Received','Return-Path']),p.name+' vollständige Header')
                for at in em.iter_attachments():
                    check(at.get_payload(decode=True)==(directory/at.get_filename()).read_bytes(),p.name+' Anhang quellenidentisch')
        combined=directory/'gesamt-pdf'/f'{case}_gesamt.pdf'
        check(combined.is_file() and len(PdfReader(combined).pages)>0,case+' Gesamt-PDF im Repository vorhanden')
        check('This test case file' not in pdftext(combined),case+' Gesamt-PDF ohne Warnhinweis')
        if ASSETS is None:
            continue
        def archive_path(name, subdir):
            for candidate in [ASSETS/name, ASSETS/subdir/name]:
                if candidate.is_file():return candidate
            raise AssertionError('Angefordertes Archiv fehlt: '+str(ASSETS/name))
        original=archive_path(f'testakte-{case}.zip','testakten')
        singles=archive_path(f'testakte-{case}-einzelpdfs.zip','testakten-einzelpdfs')
        with zipfile.ZipFile(original) as z:
            expected={p.name for p in sources}|{'README.txt',combined.name}
            check(set(z.namelist())==expected and len(z.namelist())==len(expected),case+' Original-ZIP vollständig, flach und ohne QA-Dateien')
            check(not notice_text_errors(z.read('README.txt')),case+' Original-ZIP Hinweis nach zentralem Standard')
            for p in sources:check(z.read(p.name)==p.read_bytes(),p.name+' Original-ZIP bytegleich')
            aggregate=PdfReader(io.BytesIO(z.read(combined.name)))
            exported_pdf(aggregate,case+' archiviertes Gesamt-PDF')
            for p in sources:exported_content(aggregate,p,case+' Gesamt-PDF / '+p.name)
        with zipfile.ZipFile(singles) as z:
            expected={p.stem+'.pdf' for p in sources}|{'README.txt'}
            check(set(z.namelist())==expected and len(z.namelist())==len(expected),case+' Einzel-PDF-ZIP vollständig, flach und ohne QA-Dateien')
            check(not notice_text_errors(z.read('README.txt')),case+' Einzel-PDF-ZIP Hinweis nach zentralem Standard')
            for p in sources:
                reader=PdfReader(io.BytesIO(z.read(p.stem+'.pdf')))
                exported_pdf(reader,p.name)
                exported_content(reader,p,p.name)
        if (ASSETS/combined.name).exists():
            external=PdfReader(ASSETS/combined.name)
            exported_pdf(external,case+' externes Gesamt-PDF')
            for p in sources:exported_content(external,p,case+' extern / '+p.name)


def formulas():
    anchors=[(0,'10_Kostenfortschreibung.xlsx','Kostenstand','F14',2206000),
             (0,'09_Ablauf_und_Zahlungen.xlsx','Zahlungen','G9',124860),
             (1,'24_Kreditorenabgleich.xlsx','Kreditoren','E10',3342),
             (1,'25_Projektkosten_und_Bank.xlsx','Projektkosten','D8',64480),
             (1,'55_Ergaenzungsabgleich.xlsx','Offene Posten','F20',9999),
             (1,'55_Ergaenzungsabgleich.xlsx','Bank','F15',38218.8)]
    for ci,filename,sheet,coordinate,expected in anchors:
        check(near(cell(ROOT/'testakten'/CASES[ci]/filename,sheet,coordinate),expected),filename+' fachliche Ergebniskontrolle')
    for case in CASES:
        for p in (ROOT/'testakten'/case).glob('*.xlsx'):
            form=load_workbook(p,data_only=False); values=load_workbook(p,data_only=True)
            count=0
            check(form.properties.creator=='Klotzkette',p.name+' Metadatenautor')
            check(form.properties.language=='de-DE',p.name+' Metadatensprache')
            for s in form:
                check(not s.protection.sheet,p.name+' '+s.title+' bearbeitbar')
                for row in s:
                    for c in row:
                        if c.data_type!='f':continue
                        count+=1
                        check(len(c.value)<=260 and not re.search(r'\b(?:LET|LAMBDA|MAP|REDUCE|INDIRECT|OFFSET|TABLE)\(',c.value),p.name+' '+s.title+'!'+c.coordinate+' begrenzte Formel')
                        check(not re.search(r'\$?[A-Z]+:\$?[A-Z]+',c.value),p.name+' '+c.coordinate+' begrenzter Bereich')
                        check(values[s.title][c.coordinate].data_type!='e' and values[s.title][c.coordinate].value is not None,p.name+' '+s.title+'!'+c.coordinate+' gespeicherter Rechenwert')
            check(20<=count<=300,p.name+f' {count} nachvollziehbare Formeln')
            form.close();values.close()


def cell(file,sheet,coordinate):
    book=load_workbook(file,data_only=True)
    result=book[sheet][coordinate].value
    book.close()
    return result


def mutate(source,dest,sheet,coordinate,value):
    """Nur wegwerfbare Prüfeingabe ändern, nicht das ausgelieferte artifact_tool-Modell."""
    book=load_workbook(source,read_only=True)
    index=book.sheetnames.index(sheet)+1;book.close()
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data=src.read(item.filename)
            if item.filename==f'xl/worksheets/sheet{index}.xml':
                root=ET.fromstring(data)
                target=root.find(f'.//s:c[@r="{coordinate}"]',NS)
                assert target is not None
                for child in list(target): target.remove(child)
                target.attrib.pop('t',None)
                if value is not None:
                    ET.SubElement(target,'{'+NS['s']+'}v').text=str(value)
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            out.writestr(item,data)


def recalc(source, label):
    run=QA/'office'/label
    run.mkdir(parents=True,exist_ok=True)
    output=run/source.name
    if output.exists():output.unlink()
    inputs=run/'input';inputs.mkdir(exist_ok=True)
    uncached=inputs/source.name
    # Excel-Exportwerte vollständig verwerfen, damit der Test echte Berechnung beweist.
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(uncached,'w',zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data=src.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                root=ET.fromstring(data)
                for c in root.findall('.//s:c',NS):
                    if c.find('s:f',NS) is not None:
                        for v in c.findall('s:v',NS):c.remove(v)
                        c.attrib.pop('t',None)
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            elif item.filename=='xl/workbook.xml':
                root=ET.fromstring(data)
                calc=root.find('s:calcPr',NS)
                if calc is None:calc=ET.SubElement(root,'{'+NS['s']+'}calcPr')
                calc.set('calcMode','auto');calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1')
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            out.writestr(item,data)
    cmd=[office_binary(),'-env:UserInstallation='+ (run/'profile').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(run),str(uncached)]
    p=subprocess.run(cmd,text=True,capture_output=True,timeout=150)
    (run/'office.log').write_text(p.stdout+p.stderr,encoding='utf-8')
    check(p.returncode==0 and output.exists(),label+' native LibreOffice-Neuberechnung')
    return output


def office():
    check(bool(office_binary()),'Office-Laufzeit konfiguriert')
    for case in CASES:
        for source in (ROOT/'testakten'/case).glob('*.xlsx'):
            native=recalc(source,source.stem+'-baseline')
            f=load_workbook(source,data_only=False);a=load_workbook(source,data_only=True);b=load_workbook(native,data_only=True)
            for s in f:
                for row in s:
                    for c in row:
                        if c.data_type!='f':continue
                        av,bv=a[s.title][c.coordinate].value,b[s.title][c.coordinate].value
                        check(near(av,bv) if isinstance(av,(int,float)) else av==bv,source.name+' '+s.title+'!'+c.coordinate+' Office gleich Export')
            f.close();a.close();b.close()
    mutations=[
        (0,'10_Kostenfortschreibung.xlsx','Belege','D6',38100,'Kostenstand','C14',1200100),
        (0,'10_Kostenfortschreibung.xlsx','Bestellungen','C6',118100,'Kostenstand','F14',2206100),
        (0,'10_Kostenfortschreibung.xlsx','Restleistungen','C6',100,'Kostenstand','F14',2206100),
        (0,'10_Kostenfortschreibung.xlsx','Restleistungen','C6',0,'Kostenstand','F14',2206000),
        (0,'10_Kostenfortschreibung.xlsx','Restleistungen','C6',None,'Kostenstand','F14','offen'),
        (0,'09_Ablauf_und_Zahlungen.xlsx','Ansätze','B6',46349,'Termine','F13',46370),
        (0,'09_Ablauf_und_Zahlungen.xlsx','Zahlungen','C8',500000,'Zahlungen','G9',24860),
        (1,'24_Kreditorenabgleich.xlsx','Ausgleich','C6',22700,'Kreditoren','E10',3442),
        (1,'24_Kreditorenabgleich.xlsx','Ausgleich','D7',0,'Kreditoren','E10',3603.8),
        (1,'25_Projektkosten_und_Bank.xlsx','Lohn August','D6',31,'Projektkosten','D8',64550),
        (1,'25_Projektkosten_und_Bank.xlsx','Bank','E13',0,'Bank','F13',80996.8),
        (1,'25_Projektkosten_und_Bank.xlsx','Fremdleistungen','D7',220.005,'Projektkosten','D8',64479.995),
        (1,'55_Ergaenzungsabgleich.xlsx','Zuordnung','D10',600,'Offene Posten','F20',9899),
        (1,'55_Ergaenzungsabgleich.xlsx','Bank','C6',2956,'Bank','E6',100),
        (1,'55_Ergaenzungsabgleich.xlsx','Belege','C20',-300,'Offene Posten','F20',9880),
    ]
    scratch=QA/'mutations';scratch.mkdir(exist_ok=True)
    for i,(ci,filename,sh,co,val,osh,oco,expected) in enumerate(mutations):
        source=ROOT/'testakten'/CASES[ci]/filename
        copy=scratch/f'mutation-{i:02d}.xlsx'
        mutate(source,copy,sh,co,val)
        actual=cell(recalc(copy,f'mutation-{i:02d}'),osh,oco)
        check(actual==expected if isinstance(expected,str) else near(actual,expected),f'Mutation {i}: {actual} statt {expected}')
        RESULTS['office_mutations'].append({'workbook':filename,'input':sh+'!'+co,'value':val,'output':osh+'!'+oco,'expected':expected,'actual':actual})


def quality():
    spec=importlib.util.spec_from_file_location('quality',ROOT/'scripts/validate-testakten-dokumentqualitaet.py')
    q=importlib.util.module_from_spec(spec);spec.loader.exec_module(q)
    # Nur die beiden beauftragten Akten, mit unveränderten zentralen Regeln.
    class Scope:
        def iterdir(self):return iter([ROOT/'testakten'/c for c in CASES])
        def rglob(self,pattern):return (p for c in CASES for p in (ROOT/'testakten'/c).rglob(pattern))
        def __truediv__(self,other):return ROOT/'testakten'/other
    q.TESTAKTEN=Scope()
    stream=io.StringIO()
    with contextlib.redirect_stdout(stream):result=q.main()
    log=stream.getvalue()
    (QA/'dokumentqualitaet.txt').write_text(log,encoding='utf-8')
    print(log)
    RESULTS['central_quality_exit']=result
    check(result==0,'Zentrale Dokumentqualität ohne Ausnahme bestanden')


def extra_integrity():
    case=CASES[1]
    journal=csvrows(case,'50_Ergaenzungsjournal.csv')
    invoices=[r for r in journal if amount(r['Netto_EUR'])>0]
    corrections=[r for r in journal if amount(r['Netto_EUR'])<0]
    check(len(invoices)==14 and len(corrections)==2,'Ergänzungsstapel 14 Rechnungen und zwei Korrekturen')
    check(len({r['Beleg'] for r in journal})==16,'Ergänzungsbelege eindeutig')
    check(sum(amount(r['Netto_EUR']) for r in journal)==19580,'Ergänzungsnetto 19.580 EUR')
    check(sum(amount(r['USt_EUR']) for r in journal)==Decimal('2200.20'),'Ausgewiesene Steuer 2.200,20 EUR')
    issuers={}
    for row in journal:
        source=ROOT/'testakten'/case/row['Datei']
        check(source.is_file() and row['Beleg'] in pdftext(source),'Einzelbeleg zu '+row['Beleg'])
        check(amount(row['Netto_EUR'])+amount(row['USt_EUR'])==amount(row['Belegbetrag_EUR']),'Belegrechnung '+row['Beleg'])
        if amount(row['Netto_EUR'])>0:
            text=pdftext(source)
            address=re.search(r'Gewerbestraße \d+, 32105 Bad Salzuflen',text).group()
            tax_number=re.search(r'Steuernummer [0-9/]+',text).group()
            issuer=(address,tax_number)
            check(issuers.setdefault(row['Kreditor'],issuer)==issuer,'Konstante Lieferantenstammdaten '+row['Kreditor'])
    bank=csvrows(case,'48_Projektkonto_Bank.csv')
    allocations=csvrows(case,'57_Zahlungszuordnung.csv')
    check(len(bank)==10 and len(allocations)==11,'Zehn Bankumsätze mit elf Rechnungszuordnungen')
    check(Decimal(50000)-sum(amount(r['Soll_EUR']) for r in bank)==Decimal('38218.80'),'Projektkonto 471109 Schluss 38.218,80 EUR')
    for row in bank:
        check(sum(amount(a['Zugeordnet_EUR']) for a in allocations if a['Bankreferenz']==row['Referenz'])==amount(row['Soll_EUR']),'Bankzuordnung '+row['Referenz'])
    opos=csvrows(case,'51_Ergaenzungs_OPOS.csv')
    for row in opos:
        check(amount(row['Rechnung_EUR'])-amount(row['Korrektur_EUR'])-amount(row['Zahlung_EUR'])==amount(row['Offen_EUR']),'Ergänzungs-OPOS '+row['Beleg'])
        check(sum(amount(a['Zugeordnet_EUR']) for a in allocations if a['Rechnung']==row['Beleg'])==amount(row['Zahlung_EUR']),'Zahlungsbezug '+row['Beleg'])
    check(sum(amount(r['Offen_EUR']) for r in opos)==9999,'Ergänzungs-OPOS 9.999 EUR')
    ledger=csvrows(case,'52_Buchungsvorschlaege.csv')
    accounts={r['Konto'] for r in csvrows(case,'53_Kontenstamm.csv')}
    for batch in {r['Satz'] for r in ledger}:
        check(sum(amount(r['Soll_EUR'])-amount(r['Haben_EUR']) for r in ledger if r['Satz']==batch)==0,'Ausgeglichener Buchungssatz '+batch)
    check(all(r['Konto'] in accounts for r in ledger),'Alle Vorschlagskonten im Übungsstamm')
    def balance(account):
        return sum(amount(r['Soll_EUR'])-amount(r['Haben_EUR']) for r in ledger if r['Konto']==account)
    check(balance('1401')==1520 and balance('1771')==-1520,'Paragraf-13b-Steuer und Vorsteuer je 1.520 EUR getrennt')
    check(balance('1400')==Decimal('2200.20'),'Vorsteuerberichtigung in Vorschlägen enthalten')
    check(balance('1201')==-Decimal('11781.20'),'Bankvorschläge entsprechen dem Projektkonto')
    check(sum(balance(a) for a in accounts if a.startswith('51'))==-9999,'Kreditorenkonten stimmen zum OPOS')
    check(sum(balance(a) for a in accounts if a.startswith('50'))==19580,'Aufwandskonten stimmen zum Nettostapel')


def workbook_imports():
    from akten_build_runtime import node_binary
    files=[str(p) for case in CASES for p in sorted((ROOT/'testakten'/case).glob('*.xlsx'))]
    result=subprocess.run([node_binary(),str(ROOT/'scripts/build-bauwirtschaft-kaufmaennisch-workbooks.mjs'),'--verify-imports',*files],capture_output=True,text=True,timeout=120)
    (QA/'xlsx-imports.txt').write_text(result.stdout+result.stderr,encoding='utf-8')
    check(result.returncode==0,'Alle ausgelieferten XLSX direkt mit Artifact Tool importierbar')


def readmes():
    spec=importlib.util.spec_from_file_location('downloads',ROOT/'scripts/validate-testakten-readme-downloads.py')
    validator=importlib.util.module_from_spec(spec);spec.loader.exec_module(validator)
    for case in CASES:
        errors=[]
        validator.validate_local_readme(case,ROOT/'testakten'/case,errors)
        check(not errors,case+' zentrale README-Prüfung'+(': '+'; '.join(errors) if errors else ''))


def render(renderer):
    check(Path(renderer).is_file(),'DOCX-Renderer bereitgestellt')
    for case in CASES:
        for p in (ROOT/'testakten'/case).glob('*.docx'):
            target=QA/case/'docx-render'/p.stem
            subprocess.run([sys.executable,renderer,str(p),'--output_dir',str(target),'--emit_pdf'],check=True,timeout=180)
            check(bool(list(target.glob('page-*.png'))),p.name+' DOCX-Seiten gerendert')


def render_pdfs():
    binary=os.environ.get('AKTEN_PDFTOPPM') or shutil.which('pdftoppm')
    check(bool(binary),'PDF-Renderer verfügbar')
    for case in CASES:
        target=QA/case/'pdf-render';target.mkdir(parents=True,exist_ok=True)
        for pattern in ('page-*.png','kontakt-*.png'):
            for old in target.glob(pattern):old.unlink()
        source=ROOT/'testakten'/case/'gesamt-pdf'/f'{case}_gesamt.pdf'
        subprocess.run([binary,'-scale-to','1150','-png',str(source),str(target/'page')],check=True,timeout=180)
        pages=sorted(target.glob('page-*.png'))
        check(len(pages)==len(PdfReader(source).pages),case+' sämtliche PDF-Seiten gerendert')
        for start in range(0,len(pages),4):
            contact=Image.new('RGB',(1420,2050),'#d8dcdd')
            draw=ImageDraw.Draw(contact)
            for i,p in enumerate(pages[start:start+4]):
                im=Image.open(p).convert('RGB');im.thumbnail((690,985))
                x=(i%2)*710+10;y=(i//2)*1025+30
                contact.paste(im,(x,y));draw.text((x,y-20),f'{case[-14:]} / {p.stem}',fill='black')
            contact.save(target/f'kontakt-{start//4+1:02d}.png')


def main():
    global ASSETS, QA
    parser=argparse.ArgumentParser()
    parser.add_argument('--assets',type=Path,help='ZIP-Abgleich ausdrücklich aktivieren; unterstützt Release-Unterordner.')
    parser.add_argument('--qa',type=Path,help='Verzeichnis für Prüfprotokolle; standardmäßig ein neues temporäres Verzeichnis.')
    parser.add_argument('--office',action='store_true')
    parser.add_argument('--imports',action='store_true',help='Direkten OOXML-Rückimport mit Artifact Tool prüfen.')
    parser.add_argument('--quality',action='store_true')
    parser.add_argument('--render-docx',metavar='RENDERER')
    parser.add_argument('--render-pdfs',action='store_true')
    args=parser.parse_args()
    ASSETS=args.assets.resolve() if args.assets else None
    QA=args.qa.resolve() if args.qa else Path(tempfile.mkdtemp(prefix='bauwirtschaft-kaufmaennisch-qa-'))
    QA.mkdir(parents=True,exist_ok=True)
    integrity();extra_integrity();sources_and_exports();formulas();readmes()
    if args.office:office()
    if args.imports:workbook_imports()
    if args.quality:quality()
    if args.render_docx:render(args.render_docx)
    if args.render_pdfs:render_pdfs()
    (QA/'pruefbericht.json').write_text(json.dumps(RESULTS,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(f'{len(RESULTS["checks"])} Prüfungen bestanden; {len(RESULTS["office_mutations"])} native Eingabeänderungen geprüft.')


if __name__=='__main__':main()
