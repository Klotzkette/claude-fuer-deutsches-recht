#!/usr/bin/env python3
"""Quellen-, Struktur-, Akten- und Releaseprüfungen für den Berichtsersteller."""

from __future__ import annotations

import argparse
from collections import Counter
from decimal import Decimal
from email import policy
from email.parser import BytesParser
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import unittest
from urllib.parse import unquote
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import yaml

from quality_lab import validate_profile
from testakte_disclaimer import NOTICE_BYTES
from testakte_einzelpdf_common import document_arcname_pairs
from testakte_zip_common import working_dump_archive_pairs


ROOT = Path(__file__).resolve().parents[1]
NAME = "insolvenzverwaltung-berichtsersteller"
PLUGIN = ROOT / NAME
CASES = (
    "insolvenzverwaltung-mainbogen-laborsoftware-wuerzburg",
    "insolvenzverwaltung-nordlicht-handels-kiel",
    "insolvenzverwaltung-moebelwerk-havelberg-regelverfahren",
)
DIST: Path | None = None


def module(filename: str):
    path = ROOT / "scripts" / filename
    spec = importlib.util.spec_from_file_location(path.stem.replace("-", "_"), path)
    result = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(result)
    return result


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def docx_text(path: Path) -> str:
    document = Document(path)
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


class BerichtserstellerTest(unittest.TestCase):
    def test_manifests_marketplace_and_skill_contract(self):
        for relative in ("plugin.json", ".claude-plugin/plugin.json", ".codex-plugin/plugin.json"):
            manifest = json.loads((PLUGIN / relative).read_text())
            self.assertEqual((manifest["name"], manifest["version"]), (NAME, "1.0.0"))
        market = json.loads((ROOT / ".claude-plugin/marketplace.json").read_text())
        entries = [p for p in market["plugins"] if p["name"] == NAME]
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0]["source"], "./" + NAME)
        skills = sorted((PLUGIN / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 11)
        self.assertEqual(sum(path.parent.name == "berichtserstellung-steuern" for path in skills), 1)
        for path in skills:
            text = path.read_text(encoding="utf-8")
            front = yaml.safe_load(text.split("---", 2)[1])
            self.assertEqual(set(front), {"name", "description"})
            self.assertEqual(front["name"], path.parent.name)
            self.assertLessEqual(len(front["description"]), 360)
            self.assertNotIn("§", front["description"])
            self.assertNotIn("im Plugin", front["description"])
            self.assertEqual(re.findall(r"(?m)^## ([1-6])\. ", text), list("123456"))
            self.assertGreaterEqual(len(text.split()), 1000, path)
            for required in ("Times New Roman", "11 pt", "ausformuliert", "Trägt nicht"):
                self.assertIn(required, text)
            for target in re.findall(r"\]\(([^ )]+)", text):
                if target.startswith(("http:", "https:", "#", "mailto:")):
                    continue
                destination = (path.parent / unquote(target.split("#")[0])).resolve()
                self.assertTrue(destination.exists(), f"{path}: {target}")

    def test_prompts_are_independent_bounded_and_reviewed(self):
        workshop = PLUGIN / f"{NAME}-werkstatt.md"
        mini = PLUGIN / f"{NAME}-schnellstart.md"
        self.assertEqual(workshop.read_bytes(), workshop.with_suffix(".txt").read_bytes())
        self.assertEqual(mini.read_bytes(), mini.with_suffix(".txt").read_bytes())
        self.assertLessEqual(mini.stat().st_size, 7500)
        self.assertGreaterEqual(len(workshop.read_text().split()), 4000)
        for text in (workshop.read_text(), mini.read_text()):
            for required in ("Zwischenbericht", "Schlussbericht", "Schlussrechnung", "Schlussverzeichnis", "menschliche Freigabe"):
                self.assertIn(required, text)
        profile = json.loads((ROOT / "quality/evals" / f"{NAME}.json").read_text())
        validate_profile(profile, NAME, PLUGIN)
        self.assertGreaterEqual(len(profile["sources"]), 10)
        self.assertGreaterEqual(len(profile["cases"]), 3)

    def test_sources_are_official_and_decision_limits_are_explicit(self):
        sources = (PLUGIN / "references/rechtsquellen.md").read_text()
        self.assertIn("https://www.gesetze-im-internet.de/inso/__58.html", sources)
        self.assertIn("https://www.gesetze-im-internet.de/inso/__66.html", sources)
        for section in ("IX ZB 19/20", "IX ZR 177/15", "IX ZB 16/14", "IX ZB 35/21", "IX ZB 143/11", "IX ZB 5/24", "IX ZB 13/25", "IX ZB 45/25"):
            self.assertIn(section, sources)
        self.assertGreaterEqual(sources.count("Trägt nicht"), 8)
        self.assertNotRegex(sources, r"(?i)beck-online|juris\.de|münchkomm|nzi\s|zip\s")

    def test_mainbogen_case_reconciles_and_preserves_real_attachments(self):
        directory = ROOT / "testakten" / CASES[0]
        pairs = document_arcname_pairs(directory)
        self.assertGreaterEqual(len(pairs), 37)
        self.assertEqual(Counter(p.suffix.lower() for p, _ in pairs)[".eml"], 6)
        book = load_workbook(directory / "09_Schlussrechnung_und_Kontenabgleich.xlsx", data_only=False)
        try:
            sheet = book["Abstimmung"]
            formulae = {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if cell.data_type == "f"}
            self.assertGreaterEqual(len(formulae), 6)
        finally:
            book.close()
        receipts = sum(map(Decimal, ("118450", "356200", "1109600", "286000", "94500", "120000", "58400", "21500", "20000")))
        expenses = sum(map(Decimal, ("602480.30", "354680", "141872.45", "184335.60", "126790.20", "96480.25", "21890", "37320", "4125.60", "73318")))
        self.assertEqual(receipts, Decimal("2184650.00"))
        self.assertEqual(expenses, Decimal("1643292.40"))
        self.assertEqual(receipts - expenses, Decimal("541357.60"))
        self.assertEqual(receipts - expenses - Decimal("115927.60"), Decimal("425430.00"))
        report = docx_text(directory / "08a_Schlussbericht_Referenzfassung.docx")
        for value in ("2.184.650,00 EUR", "1.643.292,40 EUR", "541.357,60 EUR", "425.430,00 EUR", "25,2335 Prozent"):
            self.assertIn(value, report)
        for open_gate in (
            "namentliche Freigabe zur Einreichung",
            "gerichtliche Zustimmung",
            "Auszahlungsvollzug",
        ):
            self.assertIn(open_gate, report)
        self.assertIn("nicht zur Einreichung", report)
        self.assertNotIn("Rechtskraftvermerk", report)
        distribution = load_workbook(directory / "12_Schlussverzeichnis.xlsx", data_only=False)
        try:
            sheet = distribution["Schlussverzeichnis"]
            quota = distribution["Quotenblatt"]
            self.assertIn("Quotenblatt", str(sheet["G13"].value))
            self.assertIn("SUM(G2:G12)", str(sheet["G13"].value))
            self.assertEqual(quota["B9"].value, "=ROUND(B2-B8,2)")
        finally:
            distribution.close()
        for email_path in sorted((directory / "eml").glob("*.eml")):
            message = BytesParser(policy=policy.default).parsebytes(email_path.read_bytes())
            self.assertFalse(message.defects)
            self.assertTrue(str(message["From"]).endswith(".example"))
            self.assertTrue(str(message["To"]).endswith(".example"))
            self.assertTrue(list(message.iter_attachments()), email_path.name)
            for attachment in message.iter_attachments():
                source = directory / attachment.get_filename()
                self.assertTrue(source.is_file())
                self.assertEqual(attachment.get_payload(decode=True), source.read_bytes())
        combined = directory / "gesamt-pdf" / f"{directory.name}_gesamt.pdf"
        reader = PdfReader(combined)
        self.assertGreaterEqual(len(reader.pages), 55)
        combined_text = "\n".join(page.extract_text() or "" for page in reader.pages)
        self.assertNotIn("###", combined_text)
        self.assertNotIn("Rechtskraftvermerk", combined_text)
        self.assertIn("Bestandskraft", combined_text)

    def test_external_template_content_did_not_enter_repository_outputs(self):
        forbidden = ("Berke GmbH", "1542 IN 1600/11", "Bethmann Bank", "Jochen Berke", "M BE 4782")
        targets = list(PLUGIN.rglob("*.md"))
        targets += [p for p in (ROOT / "testakten" / CASES[0]).rglob("*") if p.suffix.lower() in {".md", ".csv", ".eml"}]
        targets += [p for p in (ROOT / "testakten" / CASES[0]).rglob("*.docx")]
        content = []
        for path in targets:
            content.append(docx_text(path) if path.suffix.lower() == ".docx" else path.read_text(encoding="utf-8", errors="ignore"))
        joined = "\n".join(content)
        for value in forbidden:
            self.assertNotIn(value, joined)

    def test_three_case_readmes_and_expanded_closure_chains(self):
        plugin_readme = (PLUGIN / "README.md").read_text()
        for case in CASES:
            directory = ROOT / "testakten" / case
            self.assertTrue((directory / "README.md").is_file())
            self.assertTrue((directory / "gesamt-pdf" / f"{case}_gesamt.pdf").is_file())
            self.assertIn(case, plugin_readme)
            self.assertGreaterEqual(len(document_arcname_pairs(directory)), 35)
        nordlicht = ROOT / "testakten" / CASES[1]
        for name in ("27_zwischenbericht_nr_14_31-12-2023.docx", "28_schlussrechnung_kontenabstimmung_15-03-2024.xlsx", "29_freigabevermerk_abschluss_18-03-2024.docx", "31_vollzugsprotokoll_schlussverteilung_24-04-2024.docx", "33_nachtragsverteilung_steuerzins_20-10-2024.docx"):
            self.assertTrue((nordlicht / name).is_file(), name)
        self.assertFalse((nordlicht / "Insolvenzakte_Nordlicht_Handels_GmbH.pdf").exists())
        nordlicht_readme = (nordlicht / "README.md").read_text()
        for marker in ("Führende Fassung", "23,627320 %", "nicht mehr Bestandteil"):
            self.assertIn(marker, nordlicht_readme)
        self.assertEqual(len(document_arcname_pairs(nordlicht)), 53)
        self.assertEqual(yaml.safe_load((nordlicht / "rubric.yaml").read_text())["plugin"], NAME)
        havelberg = ROOT / "testakten" / CASES[2]
        expected = ("26_zwischenbericht_31-10-2026.docx", "28_schlussverzeichnis_20-01-2027.docx", "29_freigabevermerk_abschluss_23-01-2027.docx", "31_vollzugsprotokoll_schlussverteilung_12-03-2027.docx", "34_abschlusskontrolle_schlussrechnung_20-01-2027.xlsx")
        for name in expected:
            self.assertTrue((havelberg / name).is_file(), name)
        employees = (havelberg / "06_personalliste_insolvenzgeld.csv").read_text()
        self.assertIn("31,Claudia Stern,Auszubildende Tischlerin", employees)
        self.assertIn("31,Claudia Stern,Auszubildende Tischlerin,Produktion,2024-09-01,800.00,800.00,800.00,nein,2026-05-31,ja,", employees)
        self.assertNotIn("BfA", employees)
        payment_review = (havelberg / "11_15b_zahlungsanalyse.csv").read_text()
        self.assertIn("2026-04-01,Furnier Nord GmbH,18000.00,Ratenplanrate 2", payment_review)
        self.assertIn("am 07.05.2026 schriftlich veranlasst", payment_review)
        self.assertIn("Kein de-minimis-Ausschluss allein wegen der Betragshöhe", payment_review)
        hearing_report = docx_text(havelberg / "bericht_berichtstermin_14092026.docx")
        self.assertIn("33,583266 %", hearing_report)
        self.assertNotIn("18,4 %", hearing_report)
        self.assertNotRegex(hearing_report, r"(?m)^(?:I|II|III|IV|V|VI)\.")
        self.assertEqual(yaml.safe_load((havelberg / "rubric.yaml").read_text())["plugin"], NAME)

    def test_release_packages_are_complete_and_byte_exact(self):
        if DIST is None:
            self.skipTest("--dist erforderlich")
        builder = module("build-insolvenzverwaltung-berichtsersteller-release.py")
        self.assertEqual({p.name for p in DIST.iterdir()}, builder.asset_names())
        for plugin_zip, prefix in ((f"{NAME}.zip", ""), (f"{NAME}-portable.zip", NAME + "/")):
            with zipfile.ZipFile(DIST / plugin_zip) as archive:
                self.assertIsNone(archive.testzip())
                self.assertTrue(all(".." not in Path(name).parts for name in archive.namelist()))
                for source in builder.plugin_sources():
                    relative = source.relative_to(PLUGIN).as_posix()
                    self.assertEqual(archive.read(prefix + relative), source.read_bytes())
        for case in CASES:
            directory = ROOT / "testakten" / case
            with zipfile.ZipFile(DIST / f"testakte-{case}.zip") as archive:
                self.assertIsNone(archive.testzip())
                self.assertTrue(archive.read("README.txt").startswith(NOTICE_BYTES))
                expected = working_dump_archive_pairs(directory, include_gesamt_pdf=True)
                for source, arcname in expected:
                    self.assertEqual(archive.read(arcname), source.read_bytes())
            with zipfile.ZipFile(DIST / f"testakte-{case}-einzelpdfs.zip") as archive:
                self.assertIsNone(archive.testzip())
                self.assertTrue(archive.read("README.txt").startswith(NOTICE_BYTES))
                self.assertEqual(len(archive.namelist()), len(document_arcname_pairs(directory)) + 1)
        lines = (DIST / "checksums-sha256.txt").read_text().splitlines()
        self.assertEqual(len(lines), len(builder.asset_names()) - 1)
        for line in lines:
            digest, name = line.split("  ", 1)
            self.assertEqual(sha(DIST / name), digest)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dist", type=Path)
    args, remaining = parser.parse_known_args()
    DIST = args.dist.resolve() if args.dist else None
    unittest.main(argv=[__file__, *remaining])
