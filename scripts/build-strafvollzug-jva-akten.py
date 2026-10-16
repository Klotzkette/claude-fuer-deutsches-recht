#!/usr/bin/env python3
"""Vier fiktive Vollzugsakten aus kanonischen Daten reproduzierbar setzen."""
from pathlib import Path
import argparse, datetime as dt, json
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape
from testakte_disclaimer import NOTICE_MARKDOWN

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'scripts/data/strafvollzug-jva-akten.json'
FILES={
 'auftrag':'01_Auftrag_und_Verfahrensstand.docx',
 'antrag':'02_Antrag_an_die_JVA.docx','bescheid':'03_Bescheid.pdf',
 'bekanntgabe':'04_Bekanntgabevermerk.pdf','hausordnung':'05_Hausordnung_Auszug.pdf',
 'beleg1':'06_Sachbeleg.pdf','beleg2':'07_Ergaenzende_Auskunft.pdf',
 'gespraech':'08_Gespraechsvermerk.docx','gericht':'09_Gerichtsantrag_Entwurf.docx',
 'mail1':'10_Kontaktaufnahme.eml','mail2':'11_Auskunft_der_JVA.eml',
 'mail3':'12_Externe_Auskunft.eml','mail4':'13_Rueckfrage.eml',
 'excel':'14_Chronologie_Belege_Termine.xlsx','mail5':'15_Nachricht_zum_Sachstand.eml',
 'plan':'16_Vollzugsplan_Auszug.pdf'}

def seed():
    cases=[]
    cases.append(dict(
      slug='jva-berlin-besuch-ernaehrung',title='Yusuf Bender Besuch und Ernährung in Berlin',
      person='Yusuf Bender',born='12.02.1989',prison='Justizvollzugsanstalt Berlin Lindenbogen',
      address='Lindenbogen 46, 13589 Berlin',domain='jva-lindenbogen.example',record='LB-2048/26',
      lawyer='Rechtsanwältin Felicitas Abendroth',lawaddress='Havelsteg 12, 10589 Berlin',lawemail='f.abendroth@kanzlei-abendroth.example',
      contact='Dr. Hanno Winter',contactrole='Leiter der Vollzugsabteilung',
      external='Nele Bender',extmail='nele.bender@postfach.example',extaddress='Falkenwinkel 9, 14612 Falkensee',
      admission='12.05.2025',release='11.05.2027',application_date='21.09.2026',decision_date='02.10.2026',service_date='05.10.2026',
      subject='Kinderbesuch am Samstag und religiös begründete Verpflegung',
      profile=[
       'Yusuf Bender verbüßt als erwachsener Gefangener eine rechtskräftige Freiheitsstrafe. Eine Untersuchungshaft oder ein Jugendvollzug liegt nicht vor. Er arbeitet seit Februar 2026 in der Verpackungsgruppe. Seine beiden Kinder Leni und Ben sind sieben und zehn Jahre alt und leben bei ihrer Mutter Nele Bender. Die Eltern sind getrennt, stimmen die Kontakte aber gemeinsam ab.',
       'Herr Bender hat einen längeren Familienbesuch für Samstag, den 24.10.2026, und eine verlässliche Speisenauswahl beantragt. Er erläutert seine persönlichen religiösen Essensregeln und verlangt keine pauschale Behandlung aufgrund seines Namens oder einer vermuteten Gruppenzugehörigkeit. Die JVA hat zusätzliche Zeit und Samstagszugang abgelehnt; bei der Verpflegung verweist sie auf vegetarische Kost.',
       'Bearbeitungsstand ist der 16.10.2026. Der Bescheid wurde Herrn Bender am 05.10.2026 persönlich ausgehändigt. Die anwaltliche Vollmacht für diesen Vollzugsvorgang liegt in der Kanzlei vor. Der vorhandene gerichtliche Antrag ist ein noch nicht eingereichter Parteientwurf. Die Unterlagen sollen je nach Auftrag aus Sicht des Gefangenen, der Anstalt oder des Gerichts aufbereitet werden; tatsächliche und rechtliche Gegenargumente sind getrennt festzuhalten.'
      ],
      request=[
       'Sehr geehrter Herr Dr. Winter, ich beantrage für den 24.10.2026 von 10.00 bis 12.00 Uhr einen Besuch meiner Kinder Leni und Ben zusammen mit ihrer Mutter Nele Bender in einem für Kinder geeigneten Besuchsraum. Hilfsweise bitte ich um einen anderen Samstagstermin im November mit wenigstens neunzig Minuten zusammenhängender Besuchszeit. Beide Kinder möchten mir ihre Schulhefte zeigen. Die Hefte können vorab kontrolliert werden; eigenes Spielzeug oder Lebensmittel sollen nicht mitgebracht werden.',
       'Frau Bender arbeitet montags bis freitags von 8.00 bis 16.30 Uhr. Die Kinder haben an Schultagen Unterricht und anschließend Betreuung. Der bislang angebotene Mittwochstermin um 13.00 Uhr setzt Schulbefreiungen und einen unbezahlten halben Urlaubstag voraus. Die Fahrt dauert nach Auskunft von Frau Bender etwa fünfzig Minuten je Strecke. Der letzte gemeinsame Besuch war am 12.08.2026 für sechzig Minuten. Ich beantrage keine unbeaufsichtigte Begegnung. Eine Sichtaufsicht und die üblichen Einlasskontrollen akzeptiere ich.',
       'Außerdem beantrage ich eine verlässlich erkennbare Verpflegung ohne Schweinefleisch, Gelatine aus Schwein und Speisen mit zugesetztem Alkohol. Diese Regeln sind Teil meiner persönlichen religiösen Lebensführung. Eine vegetarische Hauptspeise nehme ich gern an, soweit auch ihre Zutaten dazu passen. Ich verlange weder eine bestimmte Marke noch täglich eine gesondert gekochte Mahlzeit. Mir genügt ein geeignetes Ersatzgericht oder eine einfache Zusammenstellung aus vorhandenen Zutaten.',
       'Am 14.09. und 17.09.2026 konnte mir am Ausgabeschalter niemand sagen, ob Nachspeise und Soße Gelatine beziehungsweise zugesetzten Alkohol enthielten. Ich habe die betreffenden Bestandteile weggelassen und das übrige Essen gegessen. Bitte teilen Sie mit, wie ich Zutaten künftig zuverlässig erfahren kann und welche Alternative es bei einem unklaren Bestandteil gibt. Ein Gespräch mit Küche und Abteilung kann die praktische Lösung erleichtern.',
       'Bitte entscheiden Sie über die Besuchszeit, den Samstagstermin und die einzelnen Ernährungsbestandteile gesondert. Falls der 24.10.2026 nicht möglich ist, bitte ich um konkrete Alternativtermine und um eine Erklärung, weshalb eine zusätzliche Aufsicht oder ein anderer Raum ausscheidet. Mit freundlichen Grüßen, Yusuf Bender.'
      ],
      decision=[
       '1. Entscheidung',
       'Der Antrag auf einen zweistündigen Kinderbesuch am 24.10.2026 und auf einen längeren Samstagstermin im November wird abgelehnt. Es verbleibt bei einem begleiteten Familienbesuch von sechzig Minuten während der allgemeinen Besuchszeiten. Ein Termin am Mittwoch, dem 28.10.2026, um 13.00 Uhr kann nach Bestätigung reserviert werden.',
       'Dem Antrag auf eine vegetarische Hauptmahlzeit wird entsprochen. Eine darüber hinausgehende individuelle Speisenkontrolle oder schriftliche Zutatenbestätigung für sämtliche Komponenten wird abgelehnt. Der Dienstplan der Küche sieht keine zusätzliche Ausgabezeit für Einzelberatungen vor.',
       '2. Gründe',
       'Am 24.10.2026 steht der größere Familienraum wegen einer angekündigten Wartung der Gegensprechanlage nicht zur Verfügung. Für die Samstagsbesuche sind bereits sechs Familien vorgemerkt. Eine zusätzliche zweistündige Belegung würde die für Einlass, Sichtaufsicht und Raumwechsel eingeplanten Kräfte länger binden. Der allgemeine Besuchsraum ist für wechselnde Gruppen organisiert und lässt sich nicht ohne Anpassung des Ablaufs zwei Stunden durchgehend sperren.',
       'Bei Herrn Bender wurde am 19.02.2026 im Rahmen einer Rückkehrkontrolle ein nicht zugelassener Datenspeicher aufgefunden. Der Vorgang wurde abgeschlossen; eine Beteiligung von Frau Bender oder der Kinder ist nicht festgestellt. Gleichwohl soll der Aufsichtsstandard bei Besuchen nicht abgesenkt werden. Der Antrag lässt nicht erkennen, weshalb sich die Familie nicht auf die allgemeinen Zeiten einstellen kann. Eine Arbeitgeberbescheinigung lag bei Fertigung dieses Bescheids nicht vor.',
       'Die vegetarische Verpflegung wird täglich angeboten. Eine getrennte Fertigung für weitere persönliche Zutatenwünsche würde nach Auskunft der Küche zusätzliche Arbeitsschritte und Dokumentation erfordern. Die Anstalt geht davon aus, dass mit der vegetarischen Hauptspeise eine zumutbare Auswahl eröffnet wird. Die beiden konkret genannten Ausgabetage konnten anhand der Ausgabeübersicht nicht abschließend nachvollzogen werden.',
       '3. Weiterer Ablauf',
       'Die angebotene Besuchsreservierung wird bis zum 19.10.2026 gehalten. Angaben zu einer anderen praktikablen Besuchszeit oder konkret bezeichneten Zutaten können der Abteilung vorgelegt werden. Der vorliegende Bescheid beendet die Entscheidung über den Antrag vom 21.09.2026; eine erneute Eingabe ersetzt keinen gerichtlichen Rechtsbehelf.'
      ],
      house=[
       'Auszug aus der Hausordnung der JVA Berlin Lindenbogen, interne Fassung vom 01.07.2026. Die Besuchsorganisation wird durch die Besuchsabteilung geführt. Regeltermine finden montags, mittwochs und freitags zwischen 12.00 und 16.00 Uhr statt. Samstags werden besondere Familienbesuche nach vorheriger Abstimmung durchgeführt. Eine automatische Zusage folgt aus einer Terminanfrage nicht.',
       'Besuchspersonen melden sich dreißig Minuten vor Beginn am Haupteingang. Mitgebrachte Unterlagen und Gegenstände werden vor Übergabe geprüft. Kinder dürfen altersgerechte Schulunterlagen vorzeigen; deren Übergabe bedarf der Zustimmung des Besuchsdienstes. Eine Abweichung vom vorgesehenen Raum oder Zeitfenster ist mit der Vollzugsabteilung abzustimmen.',
       'Für die Gemeinschaftsverpflegung gibt es eine täglich ausgewiesene vegetarische Hauptspeise. Fragen zu Zutaten können schriftlich über die Abteilung an die Küche gerichtet werden. Bei nicht geklärten Zutaten soll der Ausgabedienst die Küche kontaktieren. Der Auszug beschreibt den internen Ablauf und enthält keine Entscheidung über den Antrag von Herrn Bender.'
      ],
      proof1title='Bestätigung zu Arbeitszeit und Kinderbesuch',
      proof1=[
       'Nele Bender, Falkenwinkel 9, 14612 Falkensee, an die JVA Berlin Lindenbogen. 08.10.2026. Ich bestätige, dass unsere Kinder Leni und Ben ihren Vater gemeinsam besuchen möchten. Die Schule hat noch keine Befreiung für den angebotenen Mittwoch zugesagt. Ich arbeite bei der Buchhandlung Blattwerk von Montag bis Freitag bis 16.30 Uhr; mein Arbeitgeber hat für Oktober keine weitere unbezahlte Freistellung bestätigt.',
       'Am Samstag, dem 24.10.2026, können wir zwischen 9.00 und 15.00 Uhr kommen. Auch der 07.11.2026 oder 14.11.2026 wäre möglich. Wegen der Einlasszeit und der Fahrt ist ein kurzer Besuch mitten in einem Schultag schwierig. Die Kinder fragen seit dem letzten Besuch, weshalb die gemeinsame Zeit so schnell vorbei war. Eine pädagogische oder medizinische Beurteilung nehme ich damit nicht vor.',
       'Ich unterstütze die Einlasskontrollen und werde nur die angekündigten Schulhefte mitbringen. Eine Verkürzung auf neunzig Minuten wäre für uns besser als ein Ausfall. Die beiliegende Arbeitszeitinformation aus meinem Arbeitsvertrag kann ich bei Bedarf vorlegen; sie ist dieser Nachricht noch nicht beigefügt. Nele Bender.'
      ],
      proof2title='Auskunft der Küche zu Speisebestandteilen',
      proof2=[
       'JVA Berlin Lindenbogen, Küchenleitung Ruth Salbei, Vermerk vom 09.10.2026. Zum Speiseplan der Woche vom 14.09.2026 wird berichtet: Die vegetarische Hauptspeise am Montag bestand aus Kartoffeln und Gemüsebratling. Die Nachspeise wurde als zugekaufter Fruchtquark ausgegeben. Die Artikelnummer ist im Lagerbuch vermerkt, die damalige Verpackung jedoch nicht mehr vorhanden.',
       'Am Donnerstag wurde eine vegetarische Pilzsoße ausgegeben. Das Rezeptblatt führt Gemüsefond auf. Ob an diesem Tag ein anderer Fond verwendet wurde, lässt sich aus der Tagesliste nicht ablesen. Die Küche kann eine Zutatenübersicht für die derzeit verwendeten Standardprodukte erstellen; für jeden Ausgabetag wird bislang keine vollständige Zutatenkopie aufbewahrt.',
       'Eine einfache Alternative aus Brot, Käse, Obst und Rohkost ist mit Voranmeldung möglich. Dafür müsste die Abteilung der Küche rechtzeitig mitteilen, welche Bestandteile Herr Bender meidet. Ein eigener Kochdurchgang wäre hierdurch nicht zwingend erforderlich. Eine dauerhafte organisatorische Festlegung wurde noch nicht getroffen.'
      ],
      conversation=[
       'Telefonvermerk der Kanzlei vom 12.10.2026, Gespräch Felicitas Abendroth mit Nele Bender, 14.10 bis 14.28 Uhr. Frau Bender bestätigt ihre Nachricht vom 08.10.2026. Sie bittet darum, die Kinder nicht ohne Not zu ihren Gefühlen gegenüber dem Vollzug zu befragen. Sie könne die Fahrt an einem Samstag organisieren und akzeptiere eine Sichtaufsicht.',
       'Frau Bender erinnert sich an ein Telefonat mit dem Besuchsdienst, in dem von einem kleineren Raum gesprochen worden sei. Einen Namen und ein Datum kann sie nicht sicher nennen. Die Kanzlei hat daraus keine Raumzusage abgeleitet. Der schriftliche Bescheid spricht nur über den größeren Familienraum und den allgemeinen Besuchsraum.',
       'Herr Bender hat im anwaltlichen Gespräch am selben Tag erklärt, dass er Zutateninformationen und eine einfache Ersatzkomponente akzeptiert. Seine Beschwerde betrifft nach seiner Darstellung keine Pflicht zu einer bestimmten Speisezubereitung. Die aktuelle Küchenauskunft war der Kanzlei zu diesem Zeitpunkt noch nicht von der Abteilung bestätigt.'
      ],
      relief='die Anstalt unter Aufhebung der ablehnenden Teile des Bescheids vom 02.10.2026 zu verpflichten, über einen längeren Kinderbesuch an einem der konkret angebotenen Samstagstermine und über eine verlässlich geeignete Verpflegung erneut unter Berücksichtigung der individuellen Umstände zu entscheiden.',
      arguments=[
       'Der Besuchsantrag betrifft die tatsächliche Erreichbarkeit der Kinder und die Dauer ihrer Begegnung mit dem Vater. Die bloße Verweisung auf einen Mittwochmittag berücksichtigt die inzwischen vorgelegten Schul- und Arbeitsumstände nicht. Die Antragstellung zielt weder auf unbeaufsichtigten Besuch noch auf den Wegfall von Kontrollen. Die Anstalt muss deshalb nicht zwischen vollständigem Kontrollverzicht und vollständiger Ablehnung wählen.',
       'Die Reparatur des größeren Familienraums und die begrenzten Aufsichtskapazitäten sind nachvollziehbare Belange. Offen bleiben jedoch die Nutzung des kleineren Raums, eine Verkürzung auf neunzig Minuten und die beiden angebotenen Novembertermine. Der Datenspeichervorgang vom Februar kann bei der Sicherheitsprüfung berücksichtigt werden. Eine Verbindung dieses Vorgangs zu den Kindern oder ihrer Mutter wird im Bescheid ausdrücklich nicht festgestellt.',
       'Die Küche beschreibt am 09.10.2026 eine Alternative aus vorhandenen Komponenten und eine mögliche Zutatenübersicht. Das steht nicht notwendig im Widerspruch zur Ablehnung eines eigenen Kochdurchgangs, beantwortet aber eine andere und engere Frage. Über genau diese weniger aufwendige Lösung ist bislang nicht erkennbar entschieden worden. Die Antragstellerseite verlangt keine garantierte Sonderküche und keine pauschale Anerkennung beliebiger Wünsche.',
       'Für den gewünschten Besuch am 24.10.2026 ist zeitnahe Klärung erforderlich. Ob ein zusätzlicher vorläufiger gerichtlicher Antrag notwendig und tragfähig ist, hängt insbesondere davon ab, ob die Anstalt rechtzeitig einen brauchbaren Alternativtermin bestätigt. Der Entwurf enthält deshalb zunächst den Hauptsacheantrag; eine nicht geprüfte Eilentscheidung wird nicht als sicher dargestellt.'
      ],
      plan=[
       'Auszug aus der Fortschreibung des Vollzugsplans vom 14.09.2026 für Yusuf Bender, LB-2048/26. Herr Bender nimmt seine Arbeit regelmäßig wahr. Die Familie soll als stabiler Bezug in die weitere Vollzugsplanung einbezogen werden. Zwei gemeinsame Besuche beider Kinder sind seit Jahresbeginn dokumentiert; weitere Kontakte erfolgten telefonisch.',
       'Im Februar wurde ein nicht zugelassener Datenspeicher aufgefunden. Der Vorgang wurde mit der Abteilung besprochen. Seitdem sind keine weiteren vergleichbaren Feststellungen vermerkt. Der Plan ordnet daraus keine automatische Besuchssperre ab und enthält keine Einzelentscheidung über Dauer oder Zeitpunkt von Kinderbesuchen.',
       'Zur Verpflegung enthält die Fortschreibung den Hinweis, dass Herr Bender eine vegetarische Hauptspeise akzeptiert und Fragen zu einzelnen Zutaten schriftlich klären möchte. Eine Rückmeldung der Küche war beim Planungsgespräch noch offen. Die nächste Planbesprechung ist für den 16.11.2026 vorgesehen.'
      ],
      dates=[['2026-09-14','Vollzugsplan fortgeschrieben','plan','Tatsache aus Plan'],['2026-09-21','Antrag auf Besuch und Verpflegung','antrag','Antrag gestellt'],['2026-10-02','Teilweise Ablehnung','bescheid','Bescheiddatum'],['2026-10-05','Persönliche Aushändigung','bekanntgabe','Zugang dokumentiert'],['2026-10-08','Familie bietet weitere Samstage','beleg1','Neue Angaben'],['2026-10-09','Küche nennt einfache Ersatzkomponente','beleg2','Neue Auskunft'],['2026-10-12','Gespräch in der Kanzlei','gespraech','Erinnerung teils ungesichert'],['2026-10-19','Ende der angebotenen Reservierung','bescheid','Organisatorischer Termin'],['2026-10-24','Gewünschter Kinderbesuch','antrag','Nicht zugesagt']],
      openitems=['Arbeitszeitbeleg von Nele Bender fehlt noch.','Nutzung des kleineren Raums und Novemberkapazitäten sind ungeklärt.','Umsetzung der vorgeschlagenen Ersatzkomponente ist nicht bestätigt.']
    ))
    cases.append(dict(
      slug='jva-nrw-unterbringung-gesundheit',title='Mara Wendel Unterbringung und Bewegung in Nordrhein Westfalen',
      person='Mara Wendel',born='23.08.1974',prison='Justizvollzugsanstalt Essen Quellengrund',address='Quellengrund 28, 45329 Essen',domain='jva-quellengrund.example',record='QG-611/26',
      lawyer='Rechtsanwalt Till Severin',lawaddress='Lerchenufer 7, 45127 Essen',lawemail='t.severin@kanzlei-severin.example',contact='Beate Morgenroth',contactrole='Leiterin der Vollzugsabteilung',external='Liv Wendel',extmail='liv.wendel@postfach.example',extaddress='Wiesensteg 13, 45470 Mülheim an der Ruhr',admission='03.02.2026',release='02.08.2027',application_date='18.09.2026',decision_date='01.10.2026',service_date='05.10.2026',subject='Nutzbare Unterkunft mit Rollator sowie Sichtschutz und Hofzugang',
      profile=[
       'Mara Wendel befindet sich als erwachsene Frau im Vollzug einer rechtskräftigen Freiheitsstrafe. Sie benutzt einen von der Anstalt zugelassenen Rollator. Die vorliegende Akte betrifft die Nutzbarkeit der Unterkunft und die organisatorische Ermöglichung des Aufenthalts im Freien. Sie enthält keine Diagnose und keinen Auftrag zur medizinischen Therapieentscheidung.',
       'Frau Wendel lebt seit dem 11.09.2026 mit einer Mitgefangenen im Haftraum 2.14. Sie schildert Schwierigkeiten beim Drehen des Rollators zwischen Bett und Tisch und beim Erreichen der Toilette. Außerdem beanstandet sie fehlenden Sichtschutz. Die JVA verweist auf einen geeigneteren Raum, der vorerst anderweitig belegt sei, und auf die Möglichkeit, Unterstützung anzufordern.',
       'Der Bescheid wurde am 05.10.2026 ausgehändigt. Am Stichtag 16.10.2026 liegen ein Begehungsvermerk und eine ergänzende Auskunft zum Hofzugang vor. Ein ausformulierter gerichtlicher Antrag ist vorbereitet, aber nicht eingereicht. Die Bearbeitung soll Funktionsangaben, Beobachtungen, bestrittene Punkte und tatsächliche Abhilfemöglichkeiten auseinanderhalten.'
      ],
      request=[
       'Sehr geehrte Frau Morgenroth, ich beantrage eine Unterkunft, in der ich mit meinem zugelassenen Rollator Bett, Waschplatz und Toilette sicher erreichen kann. Außerdem bitte ich um einen wirksamen Sichtschutz für die Toilettennutzung und um eine verlässlich organisierte Möglichkeit, den Hof zu erreichen. Ich möchte möglichst selbstständig bleiben und nicht für jeden Toilettengang eine andere Person hinzurufen.',
       'Seit der Verlegung in Raum 2.14 am 11.09.2026 muss ich den Rollator am Tisch seitlich versetzen. Ein vollständiges Drehen gelingt mir dort nach meiner Erfahrung nicht. Der Abstand zwischen Bettende und Tisch ist kleiner als im bisherigen Raum. Meine Mitbewohnerin bemüht sich, ihre Sachen wegzuräumen; das löst die Enge aber nicht. Ich bitte um eine gemeinsame Besichtigung mit dem technischen Dienst.',
       'Die Toilette ist durch eine niedrige seitliche Wand teilweise abgeteilt. Vom Stuhl meiner Mitbewohnerin aus bleibt der obere Bereich einsehbar. Ein Vorhang ist nicht vorhanden. Ich wünsche einen Sichtschutz, der die Aufsicht und einen Notfallzugang ermöglicht. Wenn eine bauliche Änderung Zeit benötigt, bitte ich um eine umsetzbare Zwischenlösung.',
       'Am 15.09. und 17.09.2026 konnte ich nach meiner Erinnerung nicht am Hofgang teilnehmen, weil die Begleitung über die Stufe am Ausgang nicht verfügbar war. Ich verlange keine Bewertung meiner Beschwerden durch den allgemeinen Vollzugsdienst. Bitte klären Sie mit dem zuständigen Dienst, welche Unterstützung für den Weg organisatorisch benötigt wird und wie sie rechtzeitig bereitsteht.',
       'Bitte entscheiden Sie über die Raumgestaltung, den Sichtschutz und den Zugang zum Freien jeweils gesondert. Falls ein sofortiger Wechsel nicht möglich ist, bitte ich um Mitteilung der konkreten Zwischenmaßnahmen und eines überprüfbaren Zeitplans. Mit freundlichen Grüßen, Mara Wendel.'
      ],
      decision=[
       '1. Entscheidung','Der Antrag auf sofortige Verlegung aus Raum 2.14 wird abgelehnt. Der Antrag auf einen zusätzlichen beweglichen Sichtschutz wird ebenfalls abgelehnt. Dem Anliegen, die Wege zum Hof zu ermöglichen, wird dadurch entsprochen, dass Frau Wendel Unterstützung beim Stationsdienst anfordern kann. Eine feste tägliche Einzelbegleitung wird nicht angeordnet.',
       '2. Gründe','Der besser zugängliche Raum 1.03 ist derzeit einer anderen Gefangenen zugewiesen, deren Unterbringungsbedarf ebenfalls zu berücksichtigen ist. Eine kurzfristige Verdrängung ist nicht vorgesehen. Zwei weitere Räume werden wegen Leitungsarbeiten nicht genutzt. Eine Verlegung in eine andere Abteilung würde die derzeitige Gruppenzuordnung verändern und zusätzliche Abstimmungen erfordern.',
       'Der technische Dienst hat den bisherigen Raumplan ausgewertet und hält den Zugang zum Sanitärbereich grundsätzlich für möglich. Frau Wendel kann den Tisch nach Absprache mit der Mitbewohnerin verschieben. Eine konkrete Erprobung mit dem tatsächlich verwendeten Rollator fand vor dieser Entscheidung nicht statt. Bei Unterstützungsbedarf steht die Rufanlage zur Verfügung.',
       'Ein freistehender Sichtschirm kann den Notfallzugang behindern und birgt bei instabiler Aufstellung eine Sturzgefahr. Eine Befestigung an der Decke müsste vom technischen Dienst geprüft werden. Wegen der laufenden Arbeiten besteht hierfür gegenwärtig kein kurzfristiger Montagetermin. Die vorhandene Teilabtrennung wird bis dahin als ausreichend angesehen.',
       'Für den Hofgang ist der Stationsdienst zuständig. Die Dienststärke lässt eine zusätzliche Einzelbegleitung zu einer fest zugesagten Minute nicht täglich zu. Aus den Stationsnotizen ergibt sich für den 17.09.2026 eine verspätete Teilnahme; für den 15.09.2026 ist lediglich keine Teilnahme eingetragen. Ein dauerhaftes Vorenthalten des Hofgangs ist damit nicht festgestellt.',
       '3. Weitere Bearbeitung','Die technische Begehung am 07.10.2026 bleibt vorgesehen. Eine Änderung der Belegung oder ein tragfähiger Vorschlag für einen befestigten Sichtschutz kann Anlass für eine neue Entscheidung sein. Die Ablehnung des gegenwärtigen Antrags bleibt hiervon unberührt.'
      ],
      house=[
       'Hausordnung der JVA Essen Quellengrund, Abschnitt Unterbringung und Bewegungszeiten, Fassung vom 15.06.2026. Möbel bleiben grundsätzlich an den vorgesehenen Stellplätzen; Veränderungen sind mit dem Stationsdienst abzustimmen. Medizinisch zugelassene Hilfsmittel werden in die Unterbringungsplanung einbezogen. Verkehrsflächen und Notfallzugänge müssen erreichbar bleiben.',
       'Der Aufenthalt im Freien wird für Station 2 täglich in zwei Gruppen organisiert. Der Aufruf wird über die Station bekannt gegeben. Wenn für den Weg eine Unterstützung benötigt wird, soll dies dem Stationsdienst frühzeitig mitgeteilt werden. Ausfälle und verspätete Teilnahme werden im Stationsbuch festgehalten.',
       'Eigene Sichtschirme, zusätzliche Stoffbahnen und Veränderungen im Sanitärbereich sind vor der Verwendung prüfen zu lassen. Der Auszug trifft keine medizinische Feststellung und ersetzt keine Entscheidung über eine individuell erforderliche Hilfestellung.'
      ],
      proof1title='Begehungsvermerk zu Raum 2 14',
      proof1=[
       'Technischer Dienst, Paul Eichen, 07.10.2026, 9.20 bis 9.45 Uhr. Anwesend waren Frau Wendel, der Stationsdienst und der technische Dienst. Gemessen wurde zwischen Bettende und Tischkante eine freie Breite von 61 Zentimetern. Der verwendete Rollator ist nach der Kennzeichnung 62 Zentimeter breit. Ein sicherer Durchgang in der bisherigen Tischstellung wurde nicht erprobt, weil zunächst Gegenstände umgeräumt werden mussten.',
       'Nach einer Verschiebung des Tisches um 18 Zentimeter war ein gerader Durchgang möglich. Zum Wenden wurde der Rollator mehrfach vor- und zurückgesetzt. Frau Wendel erklärte, dies sei bei Müdigkeit schwierig. Eine Bewertung ihrer Belastbarkeit erfolgt durch den technischen Dienst nicht. Die zusätzliche Tischstellung verengt den Zugang zum Schrank der Mitbewohnerin.',
       'Die Sanitärabtrennung ist 1,35 Meter hoch. Vom Sitzplatz am Fenster besteht eine Blicklinie zum oberen Körperbereich einer Person auf der Toilette. Eine feste Schiene mit schwer entflammbarem Vorhang wäre nach erster Sichtprüfung technisch denkbar. Befestigung, Material und Notfallzugang wurden noch nicht abschließend geprüft. Ein Kostenvoranschlag ist beauftragt.'
      ],
      proof2title='Auszug aus den Stationsaufzeichnungen zum Hofzugang',
      proof2=[
       'JVA Essen Quellengrund, Stationsleitung, Auskunft vom 09.10.2026. Für den 15.09.2026 enthält das Stationsbuch bei Frau Wendel den Eintrag keine Teilnahme; ein Grund wurde nicht dokumentiert. Am 17.09.2026 ist vermerkt, dass Frau Wendel wegen fehlender unmittelbarer Unterstützung zunächst auf der Station blieb und zwanzig Minuten später zum Hof gebracht wurde.',
       'Am 22.09.2026 wurde die erste Gruppe wegen einer kurzfristigen Stationsmaßnahme verspätet aufgerufen. Frau Wendel nahm an der zweiten Gruppe teil. Für den 29.09.2026 liegt kein gesonderter Eintrag über einen Hilfewunsch vor. Daraus kann nicht abgeleitet werden, dass sicher kein Hilfewunsch geäußert wurde.',
       'Der Ausgang zum Hof hat eine Stufe. Eine mobile Rampe ist vorhanden, wird jedoch nicht an jedem Ausgang gelagert. Für den Zeitraum ab 12.10.2026 ist erprobt worden, die Rampe vor dem Aufruf bereitzustellen. Eine schriftliche Zuständigkeitsregelung für alle Schichten wurde noch nicht erlassen.'
      ],
      conversation=[
       'Anwaltlicher Gesprächsvermerk vom 12.10.2026, Till Severin mit Mara Wendel. Frau Wendel erklärt, die neue Tischstellung helfe beim geraden Durchgang, führe aber zu Streit über den Schrankzugang. Sie wolle ihrer Mitbewohnerin nichts vorwerfen. Sie könne den Rollator zeitweise selbst versetzen, wünsche aber keine dauerhafte Lösung, die täglich mehrere riskante Rangierbewegungen voraussetze.',
       'Frau Wendel berichtet, dass am 12.10.2026 die Rampe bereitgelegen habe und der Hofzugang funktioniert habe. Ob dies für jede Schicht abgesprochen sei, wisse sie nicht. Sie hält an ihrem Antrag auf eine verlässliche Organisation fest und möchte nicht nur einzelne günstig verlaufene Tage dokumentieren.',
       'Die Kanzlei hat keine Gesundheitsdiagnose erhoben. Für eine gerichtliche Bewertung soll die aktuelle funktionale Einschätzung des zuständigen medizinischen Dienstes angefordert werden, ohne nicht erforderliche Behandlungsunterlagen ungekürzt weiterzugeben. Der technische Begehungsvermerk liegt vor; ein konkreter Montageplan für den Sichtschutz fehlt.'
      ],
      relief='die Anstalt unter Aufhebung der ablehnenden Teile des Bescheids vom 01.10.2026 zu verpflichten, erneut über eine mit dem zugelassenen Rollator nutzbare Unterkunft, einen geeigneten Sichtschutz und eine verlässlich organisierte Unterstützung beim Hofzugang zu entscheiden.',
      arguments=[
       'Der technische Begehungsvermerk belegt eine konkrete Enge und eine Einsehbarkeit des Sanitärbereichs. Die frühere Beurteilung allein nach dem Raumplan erfasst die Nutzung mit dem tatsächlich vorhandenen Rollator nicht. Die Tischverschiebung verbessert den Durchgang, schafft aber einen neuen Konflikt beim Schrankzugang. Es wird daher keine bloß abstrakte Komfortverbesserung verlangt.',
       'Die anderweitige Belegung eines besser zugänglichen Raums ist ein ernstzunehmender Umstand. Die Antragstellerin beansprucht keine ungeprüfte Verdrängung einer anderen Gefangenen. Erforderlich bleibt eine nachvollziehbare Prüfung der verfügbaren Räume, der Zwischenmaßnahmen und des Zeitplans. Ein pauschaler Hinweis auf die Rufanlage beantwortet die Frage nach möglichst selbstständiger Toilettennutzung nicht.',
       'Die Sicherheitsbedenken gegen einen frei stehenden Schirm werden nicht bestritten. Der technische Dienst hält inzwischen einen befestigten Vorhang für prüfbar. Der Bescheid lässt nicht erkennen, weshalb bis zur Montage keine andere tragfähige Sichtschutzlösung möglich sein soll. Umfang und Dringlichkeit der Maßnahme müssen anhand der tatsächlichen Blicklinien und der gemeinsamen Belegung beurteilt werden.',
       'Zum Hofgang behauptet die Antragstellerin nicht, dass jede Teilnahme verweigert worden sei. Das Stationsbuch bestätigt mindestens eine Verzögerung durch fehlende Unterstützung und enthält für einen anderen Tag keine Begründung. Die erfolgreich erprobte Rampe spricht dafür, dass eine organisatorische Lösung möglich sein kann. Deren Verbindlichkeit und Zuständigkeit sind noch offen.'
      ],
      plan=[
       'Auszug aus der Vollzugsplanfortschreibung vom 10.09.2026 für Mara Wendel, QG-611/26. Frau Wendel soll an der Beschäftigungsgruppe im Erdgeschoss teilnehmen. Das zugelassene Hilfsmittel wird beim Wechsel der Unterkunft berücksichtigt. Die Verlegung nach Raum 2.14 erfolgt wegen der Leitungsarbeiten im bisherigen Bereich.',
       'Der Plan sieht vor, nach der Verlegung eine Rückmeldung zur Erreichbarkeit von Sanitärbereich, Gruppenraum und Hof einzuholen. Eine konkrete Raumbegehung und eine Bewertung der Blicklinien im Toilettenbereich waren zum Zeitpunkt der Fortschreibung noch nicht erfolgt.',
       'Die Teilnahme an Bewegung im Freien soll organisatorisch ermöglicht werden. Benötigte Unterstützung ist zwischen Station und zuständigem Fachdienst zu klären. Der Plan enthält weder eine bestimmte Diagnose noch eine Aussage darüber, ob Frau Wendel an einem einzelnen Tag ohne Hilfe sicher gehen konnte.'
      ],
      dates=[['2026-09-10','Planfortschreibung vor Verlegung','plan','Prüfung angekündigt'],['2026-09-18','Unterbringung und Hofzugang beantragt','antrag','Antrag gestellt'],['2026-10-01','Antrag teilweise abgelehnt','bescheid','Bescheiddatum'],['2026-10-05','Bescheid persönlich ausgehändigt','bekanntgabe','Zugang dokumentiert'],['2026-10-07','Raum mit Rollator besichtigt','beleg1','Maße und Beobachtungen'],['2026-10-09','Stationsaufzeichnungen ausgewertet','beleg2','Lücken ausdrücklich benannt'],['2026-10-12','Rampe erprobt und Gespräch geführt','gespraech','Einzelner gelungener Ablauf'],['2026-10-20','Angekündigter Rückruf Technik','mail5','Keine Montagezusage']],
      openitems=['Funktionsbezogene aktuelle Auskunft des medizinischen Dienstes fehlt.','Montageangebot und Ausweichräume sind nicht abschließend dokumentiert.','Verbindliche Zuständigkeit für die Rampe je Schicht ist offen.']
    ))
    cases.append(dict(
      slug='jva-bayern-bildung-computer',title='Korbinian Seidel Fernkurs und Offlinecomputer in Bayern',person='Korbinian Seidel',born='04.04.1993',prison='Justizvollzugsanstalt Regensburg Auenfeld',address='Auenfeld 31, 93055 Regensburg',domain='jva-auenfeld.example',record='AF-882/25',lawyer='Rechtsanwältin Daria Buchner',lawaddress='Wacholdersteg 5, 93047 Regensburg',lawemail='d.buchner@kanzlei-buchner.example',contact='Ulrich Tannhäuser',contactrole='Leiter der Vollzugsabteilung',external='Henriette Wald',extmail='henriette.wald@lernwerk-ost.example',extaddress='Fliederbogen 22, 90491 Nürnberg',admission='07.08.2024',release='06.02.2027',application_date='17.09.2026',decision_date='02.10.2026',service_date='05.10.2026',subject='Fortsetzung eines Fernkurses mit überprüfbarem Offlinegerät',
      profile=[
       'Korbinian Seidel verbüßt als erwachsener Mann eine rechtskräftige Freiheitsstrafe. Er hat früher in der Lagerverwaltung gearbeitet und möchte einen Fernkurs zur kaufmännischen Datenverarbeitung abschließen. Seit Juni bearbeitet er gedruckte Lernhefte. Für das nächste Modul sind Tabellenübungen mit einer bestimmten Softwaredatei vorgesehen.',
       'Er hat einen eigenen gebrauchten Laptop ohne Funkbetrieb und eine kontrollierte Übertragung der Kursdateien beantragt. Die Anstalt hat den Besitz abgelehnt und auf einen wöchentlichen betreuten Computertermin verwiesen. Die tatsächliche Verfügbarkeit und Softwareausstattung dieses Termins sind nach der Korrespondenz nicht abschließend geklärt.',
       'Stichtag ist der 16.10.2026. Die gerichtliche Antragsschrift ist als bearbeitbarer Parteientwurf vorbereitet. Weder ein Anspruch auf Internetzugang noch ein Recht auf beliebige Geräte soll unterstellt werden. Die Prüfung soll den konkreten Kursbedarf, die Kontrollmöglichkeiten, die Sicherheitsbedenken und die tatsächlich erreichbaren Alternativen zusammenführen.'
      ],
      request=[
       'Sehr geehrter Herr Tannhäuser, ich beantrage, für meinen laufenden Fernkurs einen überprüften Laptop im Haftraum benutzen zu dürfen. Das Gerät soll ohne Internetzugang betrieben werden. Funkmodule können fachgerecht entfernt oder dauerhaft unbrauchbar gemacht werden. Ich akzeptiere eine Prüfung vor Einbringung, eine dokumentierte Versiegelung und weitere angemessene Kontrollen.',
       'Der Kurs Lernwerk Ost, Teilnehmernummer LW-42618, umfasst im nächsten Abschnitt Tabellenkalkulation und die Bearbeitung lokaler Übungsdateien. Die Schule hat mitgeteilt, dass die Aufgaben ohne Netzverbindung bearbeitet werden können. Die Einsendung soll ausgedruckt oder über einen kontrollierten Datenträgerweg erfolgen; eine direkte Verbindung zu einem Lernportal beantrage ich nicht.',
       'Meine Schwester kann das vorhandene Gerät zur technischen Prüfung abgeben und die erforderlichen Kosten nach vorheriger Mitteilung übernehmen. Ich habe noch keine Einfuhr veranlasst. Bitte teilen Sie mit, ob ein bestimmtes Geräteprofil oder eine Prüfung durch einen von Ihnen benannten Betrieb verlangt wird. Eine pauschale Herstellerangabe möchte ich nicht als Sicherheitsnachweis darstellen.',
       'Falls der Besitz im Haftraum nicht genehmigt wird, beantrage ich hilfsweise eine tatsächlich nutzbare Kursalternative mit der erforderlichen Software und ausreichenden Bearbeitungszeiten. Der derzeit genannte Computerraumtermin am Donnerstag überschneidet sich nach meinem Plan mit dem Ende der Arbeitsgruppe. Ich bitte um Klärung, ob eine Freistellung möglich ist und ob die lokalen Kursdateien dort geöffnet werden können.',
       'Die nächste Einsendeaufgabe soll bis zum 30.10.2026 eingereicht werden. Eine Verlängerung habe ich angefragt, aber noch nicht zugesagt erhalten. Bitte entscheiden Sie unter Einbeziehung der Kursunterlagen und eines konkreten Prüfangebots. Mit freundlichen Grüßen, Korbinian Seidel.'
      ],
      decision=[
       '1. Entscheidung','Der Antrag auf Einbringung und Nutzung eines eigenen Laptops im Haftraum wird abgelehnt. Der Antragsteller kann sich für den betreuten Computerraumtermin am Donnerstag von 15.00 bis 16.00 Uhr anmelden. Ein gesonderter Anspruch auf zusätzliche Nutzungszeit oder auf die Installation der gewünschten Kurssoftware wird nicht anerkannt.',
       '2. Gründe','Ein eigenes Gerät schafft zusätzliche Speichermöglichkeiten und erschwert die Kontrolle von Dateien, Programmen und möglichen Kommunikationskomponenten. Die Erklärung, das Gerät werde offline betrieben, beseitigt diese Risiken nicht. Eine technische Kontrolle muss auch Veränderungen während der Nutzung und den Weg neuer Kursdateien erfassen. Für fortlaufende Einzelprüfungen steht derzeit nur begrenzte technische Unterstützung zur Verfügung.',
       'Das bislang vorgelegte Geräteblatt enthält weder eine vollständige Teileliste noch einen Nachweis über die Entfernung von Funkkomponenten. Eine verbindliche Zusage eines Prüfunternehmens liegt der Anstalt nicht vor. Es kann daher nicht vorausgesetzt werden, dass sich das konkrete Gerät mit vertretbarem Aufwand in einen kontrollierbaren Zustand versetzen lässt.',
       'Bei Herrn Seidel wurden im Juli 2025 nicht angemeldete Dateien auf einem damals zugelassenen Datenträger festgestellt. Hinweise auf eine Außenkommunikation ergaben sich nicht. Der Vorgang macht jedoch eine nachvollziehbare Kontrolle der künftig eingebrachten Dateien erforderlich. Die Anstalt will einen nicht überprüfbaren Austausch mit Dritten vermeiden.',
       'Der betreute Computerraum stellt grundsätzlich eine Bildungsalternative dar. Die nutzbaren Programme werden zentral betreut. Die Abteilung geht davon aus, dass die vorliegenden Übungen im Rahmen dieser Ausstattung bearbeitet werden können. Eine konkrete Probe mit den aktuellen Kursdateien wurde vor der Entscheidung nicht durchgeführt.',
       '3. Hinweise zum Ablauf','Der Bildungsdienst kann eine Verschiebung der Einsendefrist unterstützen. Über ein später eingereichtes vollständiges technisches Angebot kann erneut entschieden werden. Die gegenwärtige Ablehnung bleibt wirksam; die Anstalt hat den Kursanbieter nicht um eine verbindliche Fristverlängerung ersucht.'
      ],
      house=[
       'Hausordnung der JVA Regensburg Auenfeld, Auszug Bildungsgeräte und Datenträger, Fassung vom 01.08.2026. Elektronische Geräte dürfen erst nach schriftlicher Genehmigung und Eingangskontrolle eingebracht werden. Eine Zusendung ohne vorherige Abstimmung führt nicht zu einer Benutzungserlaubnis. Gerätekennung und genehmigter Nutzungsumfang werden dokumentiert.',
       'Der Computerraum wird unter Aufsicht genutzt. Fremde Dateien dürfen nur über den vorgesehenen Prüfweg geöffnet werden. Ungeprüfte Datenträger und die eigenständige Installation von Programmen sind nicht zugelassen. Der Bildungsdienst stimmt die Einsatzmöglichkeiten mit dem technischen Dienst ab.',
       'Die Anmeldung für den Donnerstagstermin erfolgt bis zum vorangehenden Montag. Arbeitseinsatz und Bildungszeiten sind mit der Abteilung zu koordinieren. Der vorliegende Auszug beschreibt interne Abläufe; er trifft keine Entscheidung darüber, ob diese Abläufe den konkreten Fernkurs tatsächlich ermöglichen.'
      ],
      proof1title='Kursbestätigung und Anforderungen an die Übungsdateien',
      proof1=[
       'Lernwerk Ost, Henriette Wald, Fliederbogen 22, 90491 Nürnberg. Bestätigung vom 07.10.2026 zu Korbinian Seidel, LW-42618. Herr Seidel ist seit dem 01.06.2026 für den Fernkurs Kaufmännische Datenverarbeitung eingeschrieben. Die nächsten Aufgaben behandeln Tabellenformeln, Sortierung und eine lokale Auswertung vorgegebener Geschäftsdaten.',
       'Das Modul kann ohne Internetverbindung bearbeitet werden. Erforderlich ist eine lokale Tabellenanwendung, die das mitgelieferte Dateiformat vollständig unterstützt. Die Übungsdateien enthalten erfundene Geschäftsdaten. Makros oder externe Datenverbindungen sind nicht Bestandteil dieses Moduls. Die Schule kann Dateien über einen von der Anstalt benannten Weg bereitstellen und Aufgaben alternativ als Ausdruck entgegennehmen.',
       'Eine wöchentliche Bearbeitungszeit von einer Stunde kann für manche Teilnehmer genügen, liegt aber unter der von uns empfohlenen Übungszeit. Der individuelle Bedarf wird damit nicht festgestellt. Die Einsendefrist ist derzeit der 30.10.2026. Eine Verlängerung bis zum 13.11.2026 ist möglich, sobald Herr Seidel sie ausdrücklich bestätigt; eine Bestätigung liegt uns noch nicht vor.'
      ],
      proof2title='Vorläufiges Angebot zur technischen Geräteprüfung',
      proof2=[
       'Werkstatt Datenklar, Gertrud Falk, Weidenrain 16, 93059 Regensburg. Angebot DK-261009 vom 09.10.2026 an die Familie Seidel. Wir können das bezeichnete Gebrauchtgerät vor der Einbringung öffnen, Bauteile dokumentieren und vorhandene Funkmodule identifizieren. Ein endgültiger Prüfbericht setzt die Übergabe des Geräts und eine Abstimmung der Anforderungen mit der JVA voraus.',
       'Die Entfernung eines austauschbaren Funkmoduls erscheint nach dem bisher übersandten Typenblatt möglich. Ob weitere Kommunikationskomponenten vorhanden sind, wurde noch nicht überprüft. Eine Versiegelung kann Veränderungen sichtbar machen, verhindert sie jedoch nicht sicher. Datenträgerkontrollen während der späteren Nutzung sind nicht Teil dieses einmaligen Angebots.',
       'Für Eingangssichtung, Teileliste und schriftlichen Bericht kalkulieren wir 185,00 EUR einschließlich Umsatzsteuer. Zusätzliche Arbeiten werden nur nach Rücksprache ausgeführt. Das Angebot erklärt das Gerät weder für uneingeschränkt sicher noch für vollzugstauglich. Die Festlegung des zulässigen Einsatzes bleibt bei der zuständigen Stelle.'
      ],
      conversation=[
       'Gesprächsvermerk vom 12.10.2026, Rechtsanwältin Daria Buchner mit Korbinian Seidel. Herr Seidel möchte den Kurs weiterführen, hält aber nicht um jeden Preis am eigenen Laptop fest. Er würde einen anstaltseigenen Rechner nutzen, wenn die Kursdateien geöffnet werden können und die Zeiten mit seiner Arbeit vereinbar sind.',
       'Zum Datenträgervorgang aus dem Jahr 2025 erklärt Herr Seidel, er habe zusätzlich private Texte gespeichert. Die Feststellung der Anstalt bestreitet er nicht. Er bittet darum, daraus keine tatsächlich nicht festgestellte Kommunikation über das Internet abzuleiten. Eine Überprüfung des damaligen Vorgangs ist nicht Gegenstand des jetzigen Auftrags.',
       'Für den 08.10.2026 hat der Bildungsdienst nach telefonischer Auskunft einen Computerraumtermin wegen Personalausfalls abgesagt. Die Kanzlei besitzt hierzu noch keinen schriftlichen Nachweis. Herr Seidel möchte die angebotene Fristverlängerung des Kursanbieters annehmen; eine von ihm unterschriebene Bestätigung wurde am Gesprächstag zur Weitergabe vorbereitet.'
      ],
      relief='die Anstalt unter Aufhebung des Bescheids vom 02.10.2026 zu verpflichten, über die beantragte Nutzung eines überprüfbaren Offlinegeräts, hilfsweise über eine tatsächlich geeignete betreute Kursalternative, erneut auf vollständiger Tatsachengrundlage zu entscheiden.',
      arguments=[
       'Die Antragstellerseite verlangt keinen unkontrollierten Rechner und keinen Internetzugang. Sie hat einen begrenzten Bildungszweck, lokale Dateien und einen kontrollierten Übertragungsweg benannt. Das nachgereichte Werkstattangebot ist noch kein abschließender Sicherheitsnachweis. Es eröffnet aber eine konkrete Prüfung, statt die Eignung allein aus einer allgemeinen Gerätebezeichnung abzuleiten.',
       'Die Anstalt benennt zutreffend Risiken zusätzlicher Speicher- und Kommunikationsmöglichkeiten. Eine bloße Zusage des Offlinebetriebs würde diese Risiken nicht beseitigen. Zu prüfen bleiben jedoch die tatsächlichen Komponenten, die Reichweite einer Versiegelung, der Kontrollaufwand im weiteren Betrieb und die Möglichkeit eines anstaltseigenen Geräts. Der Antragsteller ist bereit, seine Bitte auf ein überprüfbares Profil zu begrenzen.',
       'Der frühere Datenträgervorgang spricht für klare Kontrollen. Er belegt nach dem Bescheid keine Außenkommunikation. Die Begründung muss deshalb den konkreten Vorgang und die gegenwärtigen Kontrollvorschläge zusammen betrachten. Der Entwurf beansprucht keine automatische Genehmigung aufgrund bisheriger Bildungsbemühungen.',
       'Die hilfsweise angebotene Computerraumnutzung wurde weder mit den Kursdateien erprobt noch hinsichtlich Arbeitsüberschneidung und Ausfallvertretung geklärt. Die Kursbestätigung beschreibt einen realen Bedarf, lässt aber auch eine betreute Alternative zu. Über diese Alternative muss anhand der tatsächlichen Ausstattung und erreichbaren Zeiten entschieden werden. Die mögliche Fristverlängerung mindert Zeitdruck, ersetzt aber keine nutzbare Kursorganisation.'
      ],
      plan=[
       'Auszug aus dem Vollzugsplan vom 11.09.2026 für Korbinian Seidel, AF-882/25. Die Fortsetzung des Fernkurses wird als Bestandteil der beruflichen Vorbereitung aufgenommen. Herr Seidel hat die ersten gedruckten Aufgaben eingereicht. Der Bildungsdienst soll die Anforderungen der folgenden digitalen Module erheben.',
       'Die Beschäftigungsgruppe endet donnerstags regelmäßig um 15.15 Uhr. Eine abweichende Freistellung für Bildungszwecke wurde im Plan noch nicht konkret festgelegt. Eine Kursbescheinigung und eine technische Beschreibung des Geräts sollen nachgereicht werden.',
       'Der Plan verweist auf den Datenträgervorgang vom Juli 2025 und auf die Notwendigkeit eines kontrollierten Dateizugangs. Er enthält keine Genehmigung eines privaten Laptops. Die nächste Überprüfung der Bildungsorganisation soll nach Eingang der Kursbestätigung stattfinden.'
      ],
      dates=[['2026-09-11','Bildungsplanung fortgeschrieben','plan','Kursbedarf noch zu erheben'],['2026-09-17','Offlinegerät oder Kursalternative beantragt','antrag','Antrag gestellt'],['2026-10-02','Gerätenutzung abgelehnt','bescheid','Bescheiddatum'],['2026-10-05','Bescheid ausgehändigt','bekanntgabe','Zugang dokumentiert'],['2026-10-07','Kursanbieter erläutert Voraussetzungen','beleg1','Technischer Bedarf'],['2026-10-09','Werkstattangebot liegt vor','beleg2','Keine Sicherheitsfreigabe'],['2026-10-12','Alternativlösung besprochen','gespraech','Computerraum möglich'],['2026-10-30','Bisherige Einsendefrist','beleg1','Verlängerung noch zu bestätigen'],['2026-11-13','Angebotene verlängerte Einsendefrist','beleg1','Noch nicht vereinbart']],
      openitems=['Technisches Endgutachten zum tatsächlichen Gerät fehlt.','Kursdateien wurden im Computerraum noch nicht erprobt.','Fristverlängerung und Arbeitsfreistellung sind nicht abschließend bestätigt.']
    ))
    cases.append(dict(
      slug='jva-niedersachsen-lockerungen-familie',title='Samira Hartung Ausbildung Familie und Lockerungen in Niedersachsen',person='Samira Hartung',born='16.11.1991',prison='Justizvollzugsanstalt Hannover Birkenhain',address='Birkenhain 84, 30657 Hannover',domain='jva-birkenhain.example',record='BH-719/24',lawyer='Rechtsanwältin Amalia Kröger',lawaddress='Morgensteg 11, 30159 Hannover',lawemail='a.kroeger@kanzlei-kroeger.example',contact='Johannes Feddersen',contactrole='Leiter der Vollzugsabteilung',external='Mechthild Hartung',extmail='mechthild.hartung@postfach.example',extaddress='Lerchenrain 6, 30853 Langenhagen',admission='19.04.2024',release='18.01.2027',application_date='16.09.2026',decision_date='02.10.2026',service_date='05.10.2026',subject='Erprobung von Ausbildungskontakt und Familienkontakt vor Entlassung',
      profile=[
       'Samira Hartung ist erwachsen und verbüßt eine rechtskräftige Freiheitsstrafe. Ihre reguläre Entlassung ist für den 18.01.2027 vorgemerkt; eine frühere gerichtliche Entscheidung wird in dieser Akte nicht vorausgesetzt. Ihre achtjährige Tochter Fine lebt bei der Großmutter Mechthild Hartung in Langenhagen.',
       'Frau Hartung hat eine gestufte Erprobung von Außenkontakten beantragt. Sie möchte zunächst zu einem begleiteten Gespräch im Ausbildungsbetrieb, später zu festgelegten Zeiten an einer betrieblichen Vorbereitung teilnehmen und einen konkret begrenzten Familienkontakt erproben. Die JVA lehnt derzeit unbegleitete Lockerungen ab und verweist auf eine verspätete Rückkehr sowie unvollständige Angaben zum Tagesablauf.',
       'Der Stichtag ist der 16.10.2026. Betrieb und Großmutter haben inzwischen weitere Angaben geliefert. Der vorhandene gerichtliche Entwurf betrifft die individuelle Entscheidung über die beantragten Schritte. Er setzt weder eine automatische Lockerung wegen des Entlassungsdatums noch eine bestimmte Flucht- oder Missbrauchsprognose voraus.'
      ],
      request=[
       'Sehr geehrter Herr Feddersen, ich beantrage eine gestufte Vorbereitung meiner Entlassung. Zunächst möchte ich zu einem begleiteten Vorstellungsgespräch bei der Werkstatt Fadenanker in Langenhagen gehen. Der Betrieb hat für den 26.10.2026 von 10.00 bis 11.00 Uhr ein Gespräch angeboten. Anschließend bitte ich um Prüfung, ob eine regelmäßige betriebliche Vorbereitung an zwei Vormittagen pro Woche im Rahmen eines festgelegten Außenaufenthalts möglich ist.',
       'Ich beantrage nicht, dass sämtliche Schritte sofort unbegleitet stattfinden. Ich akzeptiere einen zunächst begleiteten Termin, festgelegte Wege, telefonische Rückmeldungen des Betriebs und eine erneute Entscheidung nach der ersten Erprobung. Der Betrieb möchte mit mir klären, ob ab Februar 2027 ein Ausbildungsplatz vereinbart werden kann. Ein unterschriebener Ausbildungsvertrag liegt noch nicht vor.',
       'Außerdem bitte ich um einen zeitlich begrenzten Kontakt mit meiner Tochter Fine im Haushalt ihrer Großmutter in Langenhagen. Meine Mutter kann die Betreuung und Rückmeldung übernehmen. Zunächst würde ich auch eine begleitete Begegnung oder einen kürzeren Termin akzeptieren. Über einen späteren unbegleiteten Aufenthalt soll erst nach gesonderter Prüfung entschieden werden.',
       'Bei der begleiteten Außenmaßnahme am 06.08.2026 kamen wir vierunddreißig Minuten später zurück. Nach meiner Erinnerung war die Straßenbahn ausgefallen; der Begleitdienst hat die Station telefonisch informiert. Ich bitte darum, den Dienstvermerk beizuziehen. Ich habe mich nicht von der Begleitung entfernt. Den genauen Ablauf möchte ich nicht allein aus meiner Erinnerung festlegen.',
       'Bitte berücksichtigen Sie die bisherigen Arbeits- und Bildungsrückmeldungen und benennen Sie konkrete Angaben, die für die Entscheidung noch fehlen. Ich bin bereit, die Planung an nachvollziehbare Bedingungen anzupassen. Mit freundlichen Grüßen, Samira Hartung.'
      ],
      decision=[
       '1. Entscheidung','Der Antrag auf regelmäßige unbegleitete Außenaufenthalte zur betrieblichen Vorbereitung und zum Familienkontakt wird abgelehnt. Ein begleiteter Erstkontakt kann nach weiterer organisatorischer Abstimmung erneut geprüft werden; eine verbindliche Zusage für den 26.10.2026 wird nicht erteilt.',
       '2. Gründe','Vor einer Erprobung muss der beabsichtigte Tagesablauf hinreichend bestimmt und die Rückkehr zuverlässig abgesichert sein. Das bisherige Schreiben des Betriebs nennt weder eine verantwortliche Person während der gesamten Anwesenheit noch konkrete Arbeitsbereiche. Für regelmäßige Termine fehlen eine verbindliche Zeitplanung und Angaben zur Wegeführung. Ein Ausbildungsinteresse allein beantwortet die vollzuglichen Prognosefragen nicht.',
       'Bei einer Außenmaßnahme am 06.08.2026 erfolgte die Rückkehr vierunddreißig Minuten nach der vorgesehenen Zeit. Die Stationsübersicht enthält hierzu den Eintrag verspätete Rückkehr. Der ausführliche Vermerk des Begleitdienstes wurde vor der vorliegenden Entscheidung nicht beigezogen. Hinzu kommt, dass Frau Hartung während einer internen Arbeitsphase im Mai eine Rückmeldung nach einem Raumwechsel versäumte.',
       'Die vorhandenen Arbeitsbewertungen sind überwiegend positiv. Sie zeigen ein verlässliches Verhalten im strukturierten Innenbereich, erlauben jedoch allein noch keine sichere Aussage zu unbegleiteten Außenaufenthalten. Für den Familienkontakt ist außerdem der Betreuungsablauf im Haushalt der Großmutter zu konkretisieren. Die Interessen der Tochter sind einzubeziehen; eine bloße allgemeine Berufung auf Familienbindung ersetzt die Planung nicht.',
       'Für einen begleiteten Ersttermin stehen in der Woche ab dem 26.10.2026 begrenzte Kräfte zur Verfügung. Die Anstalt kann zum jetzigen Zeitpunkt nicht zusagen, einen bestimmten Termin vollständig zu begleiten. Die konkrete Einsatzplanung wurde noch nicht abgeschlossen. Alternative Termine hat der Betrieb bislang nicht schriftlich bestätigt.',
       '3. Weitere Planung','Eine erneute Prüfung erfolgt nach Vorlage der Betriebs- und Familienangaben. Die nächste reguläre Planbesprechung ist am 20.11.2026 vorgesehen. Eine frühere Ergänzung bleibt möglich. Das vorgemerkte Entlassungsdatum führt nicht für sich allein zur Genehmigung der beantragten Maßnahmen.'
      ],
      house=[
       'JVA Hannover Birkenhain, Auszug aus dem Merkblatt zu Außenmaßnahmen, Fassung vom 01.08.2026. Vor einer genehmigten Außenmaßnahme werden Ziel, Zeiten, Wege und Kontaktpersonen schriftlich festgelegt. Änderungen sind mit der Anstalt abzustimmen. Eine telefonische Rückmeldung ersetzt keine Genehmigung einer eigenmächtigen Planänderung.',
       'Betriebe benennen eine verantwortliche Person und die vorgesehenen Tätigkeiten. Bei Ausfall eines Termins ist die Anstalt zu informieren. Für Familienkontakte werden Abhol- und Rückkehrregelung sowie die Beteiligung weiterer Personen vorab festgehalten. Die Anstalt entscheidet anhand des einzelnen Vorgangs über Begleitung und Bedingungen.',
       'Das Merkblatt dient der organisatorischen Vorbereitung. Es enthält weder eine positive Prognose für eine bestimmte Person noch eine Ablehnungsentscheidung. Voraussetzungen und Durchführung einer beantragten Maßnahme werden im jeweiligen Bescheid behandelt.'
      ],
      proof1title='Angebot eines betrieblichen Kennenlerntermins',
      proof1=[
       'Werkstatt Fadenanker, Inhaberin Rieke Sandner, Fadenweg 17, 30855 Langenhagen. Schreiben vom 08.10.2026. Für Frau Samira Hartung halten wir am 26.10.2026 von 10.00 bis 11.00 Uhr ein Gespräch über eine Ausbildung im Bereich Textilgestaltung bereit. Alternativ ist der 02.11.2026 von 10.00 bis 11.00 Uhr möglich. Eine Begleitperson kann am Gespräch teilnehmen.',
       'Für eine spätere Vorbereitung könnten wir dienstags und donnerstags von 9.00 bis 12.00 Uhr Aufgaben im Zuschnitt und in der Materialverwaltung anbieten. Ich selbst wäre als Ansprechpartnerin vor Ort. Zugang zu Bargeld oder alleinige Schlüsselgewalt wären für diese Erprobung nicht vorgesehen. Eine verbindliche Vereinbarung über die regelmäßige Teilnahme soll erst nach Abstimmung mit der JVA geschlossen werden.',
       'Wir können Beginn, Ende und besondere Vorkommnisse telefonisch zurückmelden. Die Anfahrt wäre ab der Straßenbahnhaltestelle Birkenhain mit Umstieg möglich. Einen von der JVA geprüften Wegeplan besitzen wir nicht. Das Schreiben enthält ein ernsthaftes Ausbildungsinteresse, aber noch keine Zusage eines Ausbildungsplatzes.'
      ],
      proof2title='Dienstvermerk zur begleiteten Maßnahme am 6 August',
      proof2=[
       'JVA Hannover Birkenhain, Begleitdienst, Vermerk der Bediensteten Karin Eschen vom 06.08.2026. Frau Hartung nahm von 13.00 bis 14.00 Uhr an einem begleiteten Gespräch bei der Schuldnerberatung teil. Die vorgesehene Rückkehrzeit war 15.00 Uhr. Wegen einer Streckenunterbrechung fuhr die geplante Straßenbahn nicht.',
       'Um 14.41 Uhr wurde die Station telefonisch über die Verzögerung informiert. Die Begleitung blieb während der Wartezeit und des Ersatzwegs bei Frau Hartung. Die Rückkehr wurde um 15.34 Uhr erfasst. Frau Hartung befolgte die Weisungen und entfernte sich nicht eigenmächtig. Eine andere Bewertung des Verhaltens in der konkreten Situation wird nicht festgehalten.',
       'Die Stationsübersicht erhielt anschließend den Kurzeintrag verspätete Rückkehr. Der vorliegende ausführliche Vermerk beschreibt Ursache und Ablauf dieser einen Maßnahme. Er trifft keine allgemeine Prognose zu künftigen unbegleiteten Außenaufenthalten.'
      ],
      conversation=[
       'Gesprächsvermerk vom 12.10.2026, Rechtsanwältin Amalia Kröger mit Samira Hartung. Frau Hartung erklärt, dass sie den Einstieg über einen begleiteten Betriebsbesuch akzeptiert. Sie möchte vermeiden, dass bis zur regulären Planbesprechung im November keine Terminabstimmung erfolgt und das Ausbildungsinteresse dadurch verloren geht.',
       'Der Familienkontakt soll nach ihrer Vorstellung zunächst im Wohnzimmer der Großmutter stattfinden. Die Tochter soll nicht allein für Rückmeldungen oder die Einhaltung von Bedingungen verantwortlich gemacht werden. Die Großmutter hat angeboten, für Rückfragen erreichbar zu sein. Eine pädagogische Stellungnahme zur konkreten Begegnung liegt noch nicht vor.',
       'Die Kanzlei hat den Begleitvermerk vom 06.08.2026 inzwischen erhalten. Der Vorfall im Mai bleibt gesondert zu prüfen; die günstige Erklärung des Augustvorgangs beseitigt nicht automatisch jede andere Prognosefrage. Frau Hartung bittet um einen Antrag auf erneute individuelle Entscheidung und nicht um die Behauptung eines bereits gesicherten Anspruchs auf sofortigen Freigang.'
      ],
      relief='die Anstalt unter Aufhebung des Bescheids vom 02.10.2026 zu verpflichten, erneut und getrennt über einen begleiteten Erstkontakt mit dem Ausbildungsbetrieb, die daran anschließende gestufte Erprobung sowie einen konkret geplanten Familienkontakt zu entscheiden.',
      arguments=[
       'Der Bescheid behandelt im Entscheidungssatz vor allem unbegleitete regelmäßige Maßnahmen. Der Antrag vom 16.09.2026 enthält jedoch ausdrücklich einen begleiteten Erstkontakt und weitere abgestufte Möglichkeiten. Eine Prüfung dieser engeren Schritte darf nicht dadurch offenbleiben, dass noch nicht sämtliche Voraussetzungen eines späteren regelmäßigen Außenaufenthalts feststehen.',
       'Der Dienstvermerk erläutert den zunächst nur verkürzt dokumentierten Augustvorgang. Die verspätete Rückkehr beruhte danach auf einer Verkehrsunterbrechung bei fortbestehender Begleitung und angekündigter Verzögerung. Diese Feststellungen müssen in die individuelle Bewertung einfließen. Sie erlauben keine automatische positive Prognose; ebenso wenig trägt der bloße Kurzeintrag eine ungünstige Bewertung ohne Betrachtung der Ursache.',
       'Die betrieblichen Angaben sind inzwischen wesentlich konkreter. Ansprechpartnerin, Tätigkeiten, Zeiten, Ausschluss bestimmter Zugriffe und zwei Ersttermine werden benannt. Ein geprüfter Wegeplan und eine verbindliche regelmäßige Vereinbarung fehlen weiterhin. Der Antrag zielt daher auf eine erneute Entscheidung unter Berücksichtigung dieser Angaben und der noch erforderlichen Bedingungen.',
       'Beim Familienkontakt sind die Interessen der Tochter, die Betreuung durch die Großmutter und die konkrete Ausgestaltung einzubeziehen. Das nahe Entlassungsdatum verstärkt die Bedeutung einer rechtzeitigen Vorbereitung, ersetzt aber nicht die individuelle Prognose. Die allgemeine Dienstplanlage kann konkrete Begleitmöglichkeiten begrenzen; sie beantwortet ohne Terminprüfung nicht abschließend, ob jeder angebotene Ersttermin ausscheidet.'
      ],
      plan=[
       'Auszug aus der Vollzugsplanfortschreibung vom 04.09.2026 für Samira Hartung, BH-719/24. Reguläres Entlassungsdatum laut Vollstreckungsübersicht: 18.01.2027. Frau Hartung arbeitet in der Textilgruppe. Die Rückmeldungen zur Arbeitsausführung und zum Umgang mit Material sind überwiegend positiv.',
       'Im Mai wurde eine unterlassene Rückmeldung nach einem internen Raumwechsel angesprochen. Frau Hartung hat den Vorgang im Gespräch eingeräumt. Die spätere Rückkehr von der begleiteten Maßnahme am 06.08.2026 wird in der Übersicht erwähnt; die Ursachen sollen vor der nächsten Lockerungsentscheidung anhand des Begleitberichts geklärt werden.',
       'Ausbildungsinteresse und Kontakt zur Tochter sollen in die Entlassungsvorbereitung einbezogen werden. Für eine gestufte Erprobung sind Betriebsangaben, Wege und Betreuungsablauf zu erheben. Der Plan enthält keine Vorabgenehmigung. Die nächste reguläre Besprechung ist für den 20.11.2026 eingetragen.'
      ],
      dates=[['2026-08-06','Begleitete Maßnahme mit Verkehrsverzögerung','beleg2','Ursache im Dienstvermerk'],['2026-09-04','Vollzugsplanung fortgeschrieben','plan','Bericht sollte beigezogen werden'],['2026-09-16','Gestufte Außenkontakte beantragt','antrag','Antrag gestellt'],['2026-10-02','Ablehnender Bescheid','bescheid','Bescheiddatum'],['2026-10-05','Persönliche Aushändigung','bekanntgabe','Zugang dokumentiert'],['2026-10-08','Betrieb konkretisiert Angebot','beleg1','Neue Angaben'],['2026-10-12','Stufen und Gegenargumente besprochen','gespraech','Parteiengespräch'],['2026-10-26','Erster angebotener Betriebstermin','beleg1','Nicht genehmigt'],['2026-11-02','Alternativer Betriebstermin','beleg1','Nicht genehmigt'],['2026-11-20','Reguläre Planbesprechung','plan','Organisatorischer Termin'],['2027-01-18','Vorgemerkte reguläre Entlassung','plan','Keine Vorverlegung unterstellt']],
      openitems=['Wegeplan und regelmäßige betriebliche Vereinbarung fehlen.','Begleitung an den zwei angebotenen Terminen ist nicht geprüft.','Konkrete Ausgestaltung des Familienkontakts und Fachauskunft sind offen.']
    ))
    refs={
        'jva-berlin-besuch-ernaehrung':['A1, A2 und A5','A2, A4 und A8','A1, A2 und A6','A1, A2, A5 und A7'],
        'jva-nrw-unterbringung-gesundheit':['A1 und A5','A2, A7 und A8','A2 und A5','A6 und A7'],
        'jva-bayern-bildung-computer':['A1, A5 und A6','A2, A4 und A6','A2, A7 und A8','A1, A2, A5, A7 und A8'],
        'jva-niedersachsen-lockerungen-familie':['A1 und A2','A2, A6 und A8','A1 und A5','A1, A2, A7 und A8']}
    for c in cases:c['argument_refs']=refs[c['slug']]
    return cases

def documents(c):
    # Rechtliche Texte kommen ausschließlich aus dem getrennt verifizierten Quellenpaket.
    law=c.get('verified_law',{})
    service=[f'{c["prison"]}, {c["record"]}. Bekanntgabevermerk vom 05.10.2026.',
       f'Der an {c["person"]} gerichtete Bescheid vom {de(c["decision_date"])} über {c["subject"].lower()} wurde am 05.10.2026 um 10.35 Uhr persönlich in vollständiger Ausfertigung ausgehändigt. Die empfangende Person bestätigte die Entgegennahme auf der Stationsliste. Ein bloßes Vorlesen anstelle der Aushändigung fand nicht statt.',
       'Die Ausfertigung umfasste Entscheidung, Gründe und den abschließenden Abschnitt zum weiteren Verfahren. Die Unterlagen wurden nicht gleichzeitig an eine Rechtsvertretung elektronisch zugestellt. Eine Kopie wurde auf Bitte der empfangenden Person am 06.10.2026 in den vorgesehenen Ausgang zur anwaltlichen Korrespondenz gegeben.',
       'Die vorliegende Niederschrift dokumentiert die Art und den Zeitpunkt der Bekanntgabe. Sie enthält keine Berechnung einer gerichtlichen Frist und keine Bestätigung eines späteren Eingangs bei Gericht. Aufgenommen durch '+{'jva-berlin-besuch-ernaehrung':'Alma Wiesner','jva-nrw-unterbringung-gesundheit':'Thorsten Birk','jva-bayern-bildung-computer':'Lena Herbst','jva-niedersachsen-lockerungen-familie':'Anja Lorbeer'}[c['slug']]+', Stationsdienst.']
    order=c['profile']+[f'Ansprechperson der Anstalt: {c["contact"]}, {c["contactrole"]}, {c["address"]}. Anwaltliche Korrespondenz: {c["lawyer"]}, {c["lawaddress"]}, {c["lawemail"]}.',
       'Der Arbeitsauftrag umfasst eine belegte Chronologie, die Abgrenzung der einzelnen Antragsgegenstände und die Prüfung des vorhandenen gerichtlichen Entwurfs. Fehlende Unterlagen sollen konkret angefordert werden. Ein Versand, eine Einreichung oder die Zusage einer Vollzugsmaßnahme ist nicht beauftragt.']
    judicial=[f'{c["lawyer"]}\n{c["lawaddress"]}\n{c["lawemail"]}',
       law.get('court','An die zuständige Strafvollstreckungskammer des Landgerichts'),
       f'16.10.2026\nIn der Vollzugssache {c["person"]}, geboren am {c["born"]}, {c["prison"]}, {c["address"]}, gegen die Leitung der {c["prison"]}. Vollzugsaktenzeichen: {c["record"]}.',
       '1. Anträge',f'Namens und in Vollmacht der antragstellenden Person wird beantragt, {c["relief"]}',
       'Es wird ferner um Beiziehung der vollständigen Vorgangsakte einschließlich der internen Stellungnahmen und der im Bescheid angesprochenen Planungsunterlagen gebeten. Soweit weitere tatsächliche Aufklärung erforderlich ist, wird um einen gerichtlichen Hinweis und Gelegenheit zur Ergänzung gebeten.',
       '2. Verfahrensstand',f'Der Antrag an die Anstalt datiert vom {de(c["application_date"])}. Hierüber wurde mit Bescheid vom {de(c["decision_date"])} entschieden. Die vollständige Ausfertigung wurde am 05.10.2026 persönlich ausgehändigt. Dies wird durch den Bekanntgabevermerk belegt. Der vorliegende Schriftsatz ist für eine Einreichung am 16.10.2026 vorbereitet; ein tatsächlich erfolgter gerichtlicher Eingang wird damit nicht behauptet. Beweis: Antrag, Anlage A1; Bescheid, Anlage A2; Bekanntgabevermerk, Anlage A3.',
       *law.get('procedure',[]),
       '3. Sachverhalt und Begründung',
       *[p+' Beleg: Anlagen '+ref+'.' for p,ref in zip(c['arguments'],c['argument_refs'])],
       *law.get('merits',[]),
       '4. Noch aufzuklärende Tatsachen',
       'Die antragstellende Person hält nicht sämtliche tatsächlichen Voraussetzungen für abschließend bewiesen. '+ ' '.join(c['openitems'])+' Die Anstalt soll Gelegenheit erhalten, ihre konkreten Sicherheits- und Organisationsbelange anhand der vollständigen Unterlagen zu erläutern. Der Antrag auf erneute Entscheidung nimmt das Ergebnis dieser Prüfung nicht vorweg.',
       '5. Anlagen',
       'A1: Antrag an die JVA. A2: angegriffener Bescheid. A3: Bekanntgabevermerk. A4: interner Hausordnungsauszug. A5: '+c['proof1title']+'. A6: '+c['proof2title']+'. A7: anwaltlicher Gesprächsvermerk. A8: Vollzugsplanauszug.',
       c['lawyer']]
    result={
      'auftrag':('Auftrag und Verfahrensstand',order),
      'antrag':('Antrag an die Vollzugsanstalt',[c['person']+'\n'+c['record']+'\n'+c['prison']+'\n'+c['address'],de(c['application_date'])+'\n'+c['subject']]+c['request']),
      'bescheid':('Bescheid zum Vollzugsantrag',[c['prison']+'\n'+c['address'],c['record']+'\n'+de(c['decision_date'])+'\nAn '+c['person']]+c['decision']+law.get('decision_basis',[])+law.get('notice',[])+[c['contact']+'\n'+c['contactrole']]),
      'bekanntgabe':('Vermerk über die persönliche Bekanntgabe',service),
      'hausordnung':('Auszug aus der internen Hausordnung',c['house']),
      'beleg1':(c['proof1title'],c['proof1']),
      'beleg2':(c['proof2title'],c['proof2']),
      'gespraech':('Anwaltlicher Gesprächsvermerk',c['conversation']),
      'gericht':('Antrag auf gerichtliche Entscheidung',[*judicial]),
      'plan':('Auszug aus dem Vollzugsplan',c['plan'])}
    return result

def de(value):
    return dt.date.fromisoformat(value).strftime('%d.%m.%Y') if '-' in value else value

def mails(c):
    return [
      ('mail1','2026-10-06',c['external'],c['extmail'],c['lawemail'],'Bitte um Unterstützung nach dem Bescheid',
       f'Sehr geehrte Damen und Herren,\n\nim Vorgang {c["person"]} ist nach unserer Information am 05.10.2026 ein Bescheid übergeben worden. Ich möchte die familiären beziehungsweise organisatorischen Angaben ergänzen und bitte darum, zunächst die genaue Entscheidung zu lesen. Mir liegt keine gerichtliche Eingangsbestätigung vor.\n\n{c["openitems"][0]} Bitte sagen Sie mir, welche Angaben Sie von mir benötigen und auf welchem sicheren Weg die Unterlagen übermittelt werden sollen. Ich möchte keine Zusage der JVA vorwegnehmen und auch keine Erinnerung als gesicherte Feststellung ausgeben.\n\nMit freundlichen Grüßen\n{c["external"]}\n{c["extaddress"]}',None),
      ('mail2','2026-10-08',c['contact'],'vollzug@'+c['domain'],c['lawemail'],'Hausordnung und Stand der Entscheidung',
       f'Sehr geehrte Damen und Herren,\n\nzu Ihrer Anfrage zum Vorgang {c["record"]} übersende ich den einschlägigen internen Hausordnungsauszug. Der Auszug beschreibt die derzeitigen organisatorischen Abläufe. Der konkrete Antrag wird durch den Bescheid entschieden; aus dem beigefügten Merkblatt folgt keine zusätzliche Einzelfallgenehmigung.\n\nErgänzende Unterlagen können Sie an die Vollzugsabteilung senden. Bitte kennzeichnen Sie, auf welchen Teil des Antrags sich die Angaben beziehen. Eine erneute Eingabe setzt die bestehende Entscheidung nicht außer Kraft und bestätigt keinen gerichtlichen Fristaufschub. Wir prüfen die nachgereichten Angaben und teilen mit, ob eine weitere Abstimmung erforderlich ist.\n\nMit freundlichen Grüßen\n{c["contact"]}\n{c["contactrole"]}\n{c["prison"]}', 'hausordnung'),
      ('mail3','2026-10-09',c['external'],c['extmail'],c['lawemail'],'Ergänzende schriftliche Angaben zum Vorgang',
       f'Sehr geehrte Damen und Herren,\n\nanbei übersende ich die inzwischen vorliegende Auskunft „{c["proof1title"]}“. Bitte berücksichtigen Sie das Datum des Dokuments: Diese Angaben lagen bei Erlass des Bescheids noch nicht vollständig vor. Die Datei ist eine Kopie der erhaltenen Auskunft und keine von mir erteilte Genehmigung.\n\nBitte teilen Sie mir mit, wenn die zuständige Stelle noch eine konkrete Bestätigung benötigt. Für eine telefonische Rückfrage bin ich nach vorheriger Nachricht erreichbar. Ich kann die beschriebenen Termine nicht allein verbindlich für die Anstalt zusagen. Ich wäre dankbar, wenn die nächsten Schritte mit den jeweils verantwortlichen Personen abgestimmt werden.\n\nMit freundlichen Grüßen\n{c["external"]}\n{c["extaddress"]}', 'beleg1'),
      ('mail4','2026-10-13',c['lawyer'],c['lawemail'],'vollzug@'+c['domain'],'Bitte um ergänzende Angaben und Aktenauszug',
       f'Sehr geehrte Damen und Herren,\n\nich vertrete {c["person"]} im Vorgang {c["record"]}. Bitte teilen Sie mit, ob die nachgereichten Angaben inzwischen in die Einzelprüfung aufgenommen wurden und welche konkreten Unterlagen noch fehlen. Für die weitere Bearbeitung benötige ich insbesondere folgende Punkte: {" ".join(c["openitems"])}\n\nBitte übersenden Sie außerdem die im Bescheid angesprochenen internen Stellungnahmen, soweit sie für die Entscheidung herangezogen wurden und zugänglich gemacht werden können. Für geschützte Angaben Dritter genügt zunächst eine geeignete bereinigte Fassung mit Erläuterung verbleibender Auslassungen. Eine Rücknahme des Antrags oder ein Verzicht auf gerichtliche Prüfung ist mit dieser Anfrage nicht verbunden.\n\nMit freundlichen Grüßen\n{c["lawyer"]}\n{c["lawaddress"]}',None),
      ('mail5','2026-10-15',c['contact'],'vollzug@'+c['domain'],c['lawemail'],'Zwischennachricht zur ergänzenden Prüfung',
       f'Sehr geehrte Damen und Herren,\n\ndie weiteren Angaben zum Vorgang {c["record"]} sind eingegangen. Sie wurden an die beteiligten Dienste weitergeleitet. Die angefragte vollständige Abstimmung ist noch nicht abgeschlossen. Wir planen eine telefonische Rückmeldung am 20.10.2026; dies ist kein zugesagter Durchführungstermin für die beantragte Maßnahme.\n\nDer Bescheid bleibt vorerst bestehen. Eine neue Sachentscheidung liegt dieser Nachricht nicht bei. Aus der angekündigten Rückmeldung folgt keine Zusage, dass die beantragten Schritte genehmigt werden. Soweit Sie eine zeitliche Dringlichkeit geltend machen, bitten wir um konkrete Benennung des betreffenden Termins und einer noch möglichen Alternative.\n\nMit freundlichen Grüßen\n{c["contact"]}\n{c["prison"]}',None)
    ]

def build(cases):
    from docx import Document
    from docx.shared import Mm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
    fonts=Path('/System/Library/Fonts/Supplemental');pdfmetrics.registerFont(TTFont('JvaTimes',str(fonts/'Times New Roman.ttf')));pdfmetrics.registerFont(TTFont('JvaTimesBold',str(fonts/'Times New Roman Bold.ttf')))
    ps=ParagraphStyle('Body',fontName='JvaTimes',fontSize=11,leading=14,spaceAfter=9)
    hs=ParagraphStyle('Heading',parent=ps,fontName='JvaTimesBold',spaceBefore=8,spaceAfter=9,keepWithNext=True)
    ts=ParagraphStyle('Title',parent=hs,fontSize=16,leading=19,spaceAfter=17)
    for c in cases:
        folder=ROOT/'testakten'/c['slug'];folder.mkdir(parents=True,exist_ok=True)
        for key,(title,paras) in documents(c).items():
            path=folder/FILES[key]
            if path.suffix=='.docx':
                d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(21);s.left_margin=s.right_margin=Mm(24)
                for n in ['Normal','Title','Heading 1']:
                    st=d.styles[n];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
                    for attr in list(st.element.rPr.rFonts.attrib):
                        if attr.endswith('Theme') or attr.endswith('theme'):del st.element.rPr.rFonts.attrib[attr]
                    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
                d.styles['Normal'].paragraph_format.space_after=Pt(8);d.styles['Normal'].paragraph_format.line_spacing=1.06
                d.styles['Title'].font.size=Pt(16);d.styles['Title'].font.bold=True
                d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(9);d.styles['Heading 1'].paragraph_format.space_after=Pt(9)
                d.add_paragraph(title,'Title')
                for text in paras:d.add_paragraph(text,'Heading 1' if len(text)<70 and text[:1].isdigit() and '. ' in text[:5] else 'Normal')
                footer=s.footer.paragraphs[0];footer.alignment=2;r=footer.add_run();r.font.name='Times New Roman';r.font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.append(fld)
                d.core_properties.author=c['lawyer'] if key in ('auftrag','gericht','gespraech') else c['person'];d.core_properties.language='de-DE';d.core_properties.title=title;d.save(path)
            else:
                story=[Paragraph(escape(title),ts)]
                for i,text in enumerate(paras):
                    style=hs if len(text)<70 and text[:1].isdigit() and '. ' in text[:5] else ps
                    if key=='bescheid' and i==len(paras)-2:style=ParagraphStyle('SignatureTogether',parent=style,keepWithNext=True)
                    story.append(Paragraph(escape(text).replace('\n','<br/>'),style))
                def page(can,doc):can.setFont('JvaTimes',9);can.drawString(54,27,c['record']);can.drawRightString(A4[0]-54,27,str(doc.page))
                SimpleDocTemplate(str(path),pagesize=A4,leftMargin=54,rightMargin=54,topMargin=47,bottomMargin=47,title=title,author=c['prison']).build(story,onFirstPage=page,onLaterPages=page)
        for key,date,sender,mailbox,to,subject,body,attachment in mails(c):
            msg=EmailMessage(policy=SMTP);msg['From']=f'{sender} <{mailbox}>';msg['To']=to;msg['Subject']=subject;msg['Date']=format_datetime(dt.datetime.fromisoformat(date+'T10:15:00').replace(tzinfo=ZoneInfo('Europe/Berlin')));msg['Message-ID']=f'<{c["slug"]}-{key}-{date}@{c["domain"]}>';msg['Content-Language']='de-DE';msg.set_content(body)
            if attachment:msg.add_attachment((folder/FILES[attachment]).read_bytes(),maintype='application',subtype='pdf',filename=FILES[attachment])
            (folder/FILES[key]).write_bytes(msg.as_bytes())
        (folder/'README.md').write_text(readme(c))
        (folder/'rubric.yaml').write_text(f'name: {c["slug"]}\nplugin: strafvollzug-jva\ndescription: "{c["title"]}: individuelle Vollzugsentscheidung mit Belegen und Parteientwurf."\nchecks:\n  - id: originale\n    check_type: working_file_count\n    min: 16\n    description: "Sechzehn native Aktenstücke liegen vor."\n  - id: bekanntgabe\n    check_type: human_review\n    description: "Bescheiddatum, Bekanntgabe, Einreichung und Organisationsfristen werden getrennt."\n  - id: gegenargumente\n    check_type: human_review\n    description: "Sicherheitsbelange, praktische Alternativen und Beleglücken werden individuell geprüft."\n  - id: rechtsweg\n    check_type: human_review\n    description: "Landesrecht und gerichtlicher Rechtsschutz werden anhand des tatsächlichen Vollzugsstatus geprüft."\n')
    print(f'{len(cases)} Akten erstellt; Excel folgt durch build-strafvollzug-jva-akten.mjs.')

def readme(c):
    return f'''# Testakte {c['title']}

Diese Arbeitsakte dient zur Prüfung einer individuellen Entscheidung im Erwachsenenstrafvollzug. Sie kann aus Sicht der betroffenen Person, der JVA oder der Strafvollstreckungskammer bearbeitet werden. Stichtag ist der 16.10.2026. Lesen Sie zunächst den Verfahrensstand, Antrag, Bescheid und Bekanntgabevermerk; vergleichen Sie danach die späteren Belege und den noch nicht eingereichten gerichtlichen Parteientwurf.

## 1. Arbeitsauftrag und Formate

Die Akte enthält sechzehn Originaldateien: vier Word-Dokumente, sechs PDFs, fünf E-Mails mit zwei echten PDF-Anhängen und eine bearbeitbare Excel-Datei. Die Excel-Datei erklärt Chronologie, Belegzuordnung und Termine. Organisatorische Termine sind keine berechneten gesetzlichen Fristen. Der gerichtliche Entwurf ist eine Parteiposition und keine Musterlösung; Sicherheitsbelange und offene Tatsachen bleiben sichtbar. Belege A1 bis A8 sind im Entwurf und im Belegblatt zugeordnet.

## 2. Fiktiver Vorgang

Alle beteiligten Personen, die JVA, Unternehmen, Anschriften und Vorgänge sind für diese Akte erfunden. Es handelt sich um rechtskräftige Erwachsenenstrafhaft, nicht Untersuchungshaft oder Jugendvollzug. Die reservierten `.example`-Adressen sind keine Versandziele. Interne Hausordnungen und Vollzugspläne sind Sachverhaltsbelege, keine amtlichen Rechtsquellen. Die Akte enthält keine echten Gesundheits- oder Gefangenendaten.

<!-- reserved-example-contacts -->

<!-- BEGIN gesamt-pdf-section (autogen) -->
## 3. Downloads

{NOTICE_MARKDOWN}

| Format | Download |
| --- | --- |
| Gesamt-PDF | [Lesefassung](gesamt-pdf/{c['slug']}_gesamt.pdf) |
| Akten-ZIP | [Originaldateien](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/strafvollzug-jva-v1.0.0/testakte-{c['slug']}.zip) |
| Einzel-PDF-ZIP | [Einzelne PDF-Dokumente](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/strafvollzug-jva-v1.0.0/testakte-{c['slug']}-einzelpdfs.zip) |

<!-- END gesamt-pdf-section (autogen) -->
'''

def main():
    p=argparse.ArgumentParser();p.add_argument('--seed',action='store_true');p.add_argument('--data-only',action='store_true');p.add_argument('--readmes-only',action='store_true');args=p.parse_args()
    if args.seed or not DATA.exists():DATA.parent.mkdir(parents=True,exist_ok=True);DATA.write_text(json.dumps(seed(),ensure_ascii=False,indent=2)+'\n')
    cases=json.loads(DATA.read_text())
    if args.readmes_only:
        for c in cases:(ROOT/'testakten'/c['slug']/'README.md').write_text(readme(c))
    elif not args.data_only:build(cases)

if __name__=='__main__':main()
