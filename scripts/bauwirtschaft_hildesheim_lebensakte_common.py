"""Ausgabehilfen für die erweiterte Hildesheimer Projektakte. Autor: Klotzkette."""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime
import importlib.util
import io
import os
from pathlib import Path
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SLUG = 'bauwirtschaft-hildesheim-lebensakte'
CASE = ROOT / 'testakten' / SLUG
BASE_CASE = ROOT / 'testakten/bauwirtschaft-neubau-achtfamilienhaus-hildesheim'
ASSETS = Path(os.environ.get('HILDESHEIM_LEBENSAKTE_ASSETS', '/tmp/hildesheim-lebensakte'))
QUALITY = ROOT / 'quality/hildesheim-lebensakte'
REF = 'SW-HI-26-08'
AUTHOR = 'Klotzkette'

# Eigener Modulkontext: Kein Generator verändert die Pfade des Ausgangsfalls.
_spec = importlib.util.spec_from_file_location('hildesheim_lebensakte_legacy', ROOT / 'scripts/bauwirtschaft_hildesheim_common.py')
_legacy = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_legacy)
_legacy.CASE = CASE
_legacy.ASSETS = ASSETS
ACTORS = deepcopy(_legacy.ACTORS)
ACTORS.update({
    'objektueberwachung': ('Konturfeld Architektur PartG mbB', 'Planerhof 12, 31134 Hildesheim', 'hans.mueller@konturfeld.example', 'Hans Müller, Objektüberwachung'),
    'polier': ('Steinwerk Hochbau GmbH', 'Werkhof 21, 31135 Hildesheim', 'georg.schmidt@steinwerk.example', 'Georg Schmidt, Polier'),
    'projektsteuerung': ('Steinbogen Wohnen GmbH', 'Wohnhof Am Steinbogen 18, 31134 Hildesheim', 'friedrich.weber@steinbogen-wohnen.example', 'Friedrich Weber, Projektsteuerung'),
    'buchhaltung': ('Steinbogen Wohnen GmbH', 'Wohnhof Am Steinbogen 18, 31134 Hildesheim', 'erna.wagner@steinbogen-wohnen.example', 'Erna Wagner, Buchhaltung'),
    'notar_option': ('Notar Dr. Johann Bauer', 'Urkundenweg 14, 31134 Hildesheim', 'kanzlei@notar-bauer.example', 'Dr. Johann Bauer, Notar'),
})
_legacy.ACTORS = ACTORS
euro = _legacy.euro
written_norms = _legacy.written_norms
setup_fonts = _legacy.setup_fonts


def _prepare(record):
    relative = Path(record['filename'])
    if relative.is_absolute() or '..' in relative.parts:
        raise ValueError('Ausgabepfad muss innerhalb der Projektakte liegen')
    (CASE / relative).parent.mkdir(parents=True, exist_ok=True)
    return record


def docx(record):
    _legacy.docx(_prepare(record))


def pdf(record):
    _legacy.pdf(_prepare(record))


def eml(record):
    _legacy.eml(_prepare(record))


def new_document(title, date, *, header=True):
    """A4-Arbeitsdokument mit Times New Roman 11 und nutzbaren Tabellenstilen."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    section.header_distance = section.footer_distance = Cm(.8)
    for style in document.styles:
        if style.type == 1:
            style.font.name = 'Times New Roman'
            style.font.size = Pt(11)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_after = Pt(7)
            style.paragraph_format.line_spacing = 1.05
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if key.endswith('Theme'):
                    del fonts.attrib[key]
            for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                fonts.set(qn('w:' + key), 'Times New Roman')
            for border in style.element.xpath('./w:pPr/w:pBdr'):
                border.getparent().remove(border)
    for name, size in [('Title', 16), ('Heading 1', 13), ('Heading 2', 11), ('Heading 3', 11)]:
        style = document.styles[name]
        style.font.size, style.font.bold = Pt(size), True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    properties = document.core_properties
    properties.author = properties.last_modified_by = AUTHOR
    properties.title, properties.language = title, 'de-DE'
    properties.created = properties.modified = datetime.fromisoformat(date)
    if header:
        paragraph = section.header.paragraphs[0]
        paragraph.text = REF + ' · Wohnhof Am Steinbogen in Hildesheim'
        for run in paragraph.runs:
            run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.text = REF + ' · Seite '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.add_paragraph(title, 'Title')
    return document


def add_table(document, rows, widths=None):
    """Vergleichbare Datensätze mit wiederholtem Kopf und mitwachsender Höhe."""
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    widths = widths or [17.4 / len(rows[0])] * len(rows[0])
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + name)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            edge.set(qn('w:' + key), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.text = str(value)
            properties = cell._tc.get_or_add_tcPr()
            margin = OxmlElement('w:tcMar')
            for name in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + name)
                edge.set(qn('w:w'), '85')
                edge.set(qn('w:type'), 'dxa')
                margin.append(edge)
            properties.append(margin)
            if index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E8EDF0')
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = index == 0
    document.add_paragraph()
    return table


def save_docx(document, relative):
    relative = Path(relative)
    if relative.is_absolute() or '..' in relative.parts:
        raise ValueError('Ungültiger Dokumentpfad')
    target = CASE / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    stream = io.BytesIO()
    document.save(stream)
    target.write_bytes(stream.getvalue())
    normalize_docx(target)
    return target


def normalize_docx(target):
    """Normalisiert Container-Zeitstempel ohne Änderungen am Dokumentinhalt."""
    content = Path(target).read_bytes()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as output:
        for name in source.namelist():
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            output.writestr(info, source.read(name))
    return Path(target)
