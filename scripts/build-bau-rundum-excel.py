#!/usr/bin/env python3
"""Erzeugt zehn individuelle Arbeitsmappen und prüft die native Neuberechnung."""

import argparse
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from bau_rundum_excel import ROOT, specifications, update_readmes

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", NS)


def print_profile(file, spec):
    with zipfile.ZipFile(file) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    workbook = ET.fromstring(parts["xl/workbook.xml"])
    names = workbook.find(f"{{{NS}}}definedNames")
    if names is None:
        names = ET.Element(f"{{{NS}}}definedNames")
        calc = workbook.find(f"{{{NS}}}calcPr")
        workbook.insert(list(workbook).index(calc) if calc is not None else len(workbook), names)
    for i, sheet in enumerate(spec["sheets"]):
        name = f"xl/worksheets/sheet{i + 1}.xml"
        root = ET.fromstring(parts[name])
        props = root.find(f"{{{NS}}}sheetPr")
        if props is None:
            props = ET.Element(f"{{{NS}}}sheetPr")
            root.insert(0, props)
        setup = props.find(f"{{{NS}}}pageSetUpPr")
        if setup is None:
            setup = ET.SubElement(props, f"{{{NS}}}pageSetUpPr")
        setup.set("fitToPage", "1")
        for tag in ("pageMargins", "pageSetup", "headerFooter"):
            node = root.find(f"{{{NS}}}{tag}")
            if node is not None:
                root.remove(node)
        # Insert print elements before drawings/tableParts in SpreadsheetML order.
        anchor = next((n for n in root if n.tag.split('}')[-1] in ("rowBreaks", "colBreaks", "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing", "tableParts", "extLst")), None)
        pos = list(root).index(anchor) if anchor is not None else len(root)
        margins = ET.Element(f"{{{NS}}}pageMargins", dict(left="0.3", right="0.3", top="0.4", bottom="0.4", header="0.15", footer="0.15"))
        paper = "8" if sum(c["width"] for c in sheet["columns"]) > 130 else "9"
        page = ET.Element(f"{{{NS}}}pageSetup", dict(paperSize=paper, orientation="landscape", fitToWidth="1", fitToHeight="0"))
        footer = ET.Element(f"{{{NS}}}headerFooter")
        ET.SubElement(footer, f"{{{NS}}}oddFooter").text = "&L&A&CSeite &P von &N&R" + spec["date"]
        for node in (margins, page, footer):
            root.insert(pos, node)
            pos += 1
        title = sheet["name"].replace("'", "''")
        for key, value in (("_xlnm.Print_Titles", f"'{title}'!$8:$8"), ("_xlnm.Print_Area", f"'{title}'!$A$1:${chr(64 + len(sheet['columns']))}${8 + len(sheet['rows'])}")):
            for n in list(names):
                if n.get("name") == key and n.get("localSheetId") == str(i):
                    names.remove(n)
            ET.SubElement(names, f"{{{NS}}}definedName", name=key, localSheetId=str(i)).text = value
        parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    parts["xl/workbook.xml"] = ET.tostring(workbook, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(file, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def native_convert(soffice, source, target, profile, extension):
    target.mkdir(parents=True, exist_ok=True)
    result = subprocess.run([str(soffice), f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", extension, "--outdir", str(target), str(source)], check=True, capture_output=True, text=True, timeout=120)
    output = target / f"{source.stem}.{extension}"
    if not output.is_file():
        raise RuntimeError(result.stdout + result.stderr)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    parser.add_argument("--runtime", type=Path, default=Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies")
    parser.add_argument("--qa-root", type=Path, default=Path("/tmp/bau-rundum-excel-qa"))
    parser.add_argument("--case")
    parser.add_argument("--install", action="store_true")
    args = parser.parse_args()
    from openpyxl import load_workbook
    output, qa = args.output.resolve(), args.qa_root.resolve()
    output.mkdir(parents=True, exist_ok=True)
    qa.mkdir(parents=True, exist_ok=True)
    node = args.runtime / "node/bin/node"
    soffice = Path(os.environ.get("SOFFICE", args.runtime / "bin/override/soffice"))
    with tempfile.TemporaryDirectory(prefix="bau-excel-") as temp:
        work = Path(temp)
        (work / "node_modules").symlink_to(args.runtime / "node/node_modules", target_is_directory=True)
        builder = work / "builder.mjs"
        shutil.copyfile(ROOT / "scripts/build-bau-rundum-excel.mjs", builder)
        subprocess.run([str(node), str(builder), str(ROOT), str(output), str(qa), *([args.case] if args.case else [])], check=True, timeout=600)
        for spec in specifications():
            if args.case and spec["slug"] != args.case:
                continue
            source = output / spec["slug"] / spec["filename"]
            print_profile(source, spec)
            review = qa / spec["slug"]
            final = native_convert(soffice, source, work / spec["slug"], work / "office-profile", "xlsx")
            shutil.copyfile(final, source)
            mutation = native_convert(soffice, review / "mutation.xlsx", work / spec["slug"], work / "office-profile", "xlsx")
            mutated = load_workbook(mutation, data_only=True)
            m = spec["mutation"]
            assert abs(mutated[m["outputSheet"]][m["outputCell"]].value - m["expected"]) < 0.00001
            mutated.close()
            values = load_workbook(source, data_only=True)
            for control in spec["controls"]:
                actual = values[control["sheet"]][control["cell"]].value
                expected = control["value"]
                assert abs(actual - expected) < 0.00001 if isinstance(expected, (int, float)) else actual == expected, (spec["slug"], control, actual)
            values.close()
            native_convert(soffice, source, review / "pdf", work / "office-profile", "pdf")
            if args.install:
                shutil.copyfile(source, ROOT / "testakten" / spec["slug"] / spec["filename"])
            print(f"{spec['slug']}: native Neuberechnung und Mutation bestanden.", flush=True)
    if args.install:
        update_readmes()


if __name__ == "__main__":
    main()
