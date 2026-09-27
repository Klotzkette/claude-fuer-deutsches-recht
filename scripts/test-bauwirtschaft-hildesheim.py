#!/usr/bin/env python3
"""Unabhängige Aktenregression; optional Release-Pakete mit --assets DIR. Autor: Klotzkette.

Liest ausschließlich fertige Originale und Finanzdaten, importiert keine Builder
und erzeugt keine Dateien. Die UBL-Prüfung ersetzt keine XSD-/Schematronprüfung.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from functools import lru_cache
import io
import json
from pathlib import Path
import re
import sys
import unittest
import unicodedata
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from PIL import Image, ImageStat
from pypdf import PdfReader
from testakte_disclaimer import NOTICE_BYTES, NOTICE_DE, NOTICE_EN


ROOT = Path(__file__).resolve().parents[1]
SLUG = 'bauwirtschaft-neubau-achtfamilienhaus-hildesheim'
CASE = ROOT/'testakten'/SLUG
FINANCE = ROOT/'quality/hildesheim-achtfamilienhaus/finanzdaten.json'
CENT = Decimal('0.01')
FINAL_COST = Decimal('3482215.41')
FINAL_CASH = Decimal('217784.59')
BASE_COST = Decimal('3463770.41')
AREAS = [68, 75, 82, 68, 75, 82, 75, 75]
FORMATS = {'.pdf', '.docx', '.xlsx', '.eml', '.png', '.csv', '.txt', '.xml'}
PRIMARY_NUMBERS = set(range(1, 117)) | set(range(119, 123)) | set(range(130, 180))
NS = {
    'cbc': 'urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2',
    'cac': 'urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2',
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'dc': 'http://purl.org/dc/elements/1.1/',
}
ASSETS = None


def amount(value):
    return Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)


def german(value):
    return f'{amount(value):,.2f}'.replace(',', '_').replace('.', ',').replace('_', '.')


def clean(text):
    return ' '.join(text.split())


@lru_cache(maxsize=None)
def pdf_text(path):
    return clean(' '.join(page.extract_text() or '' for page in PdfReader(path).pages))


def docx_text(path):
    with zipfile.ZipFile(path) as z:
        node = ET.fromstring(z.read('word/document.xml'))
    return clean(' '.join(n.text or '' for n in node.findall('.//w:t', NS)))


def originals():
    return sorted((p for p in CASE.iterdir() if p.is_file() and re.match(r'^\d{3}_', p.name)), key=lambda p: p.name)


def email_message(path):
    return BytesParser(policy=policy.default).parsebytes(path.read_bytes())


def words(text):
    """Inhalt unabhängig von Schrift, Silbentrennung und Seitenumbrüchen."""
    text = unicodedata.normalize('NFKC', text).replace('\u00ad', '')
    text = re.sub(r'(?<=\w)-\s*\n\s*(?=\w)', '', text)
    return Counter(re.findall(r'[^\W_]+', text.casefold()))


class HildesheimCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = json.loads(FINANCE.read_text(encoding='utf-8'), parse_float=Decimal)
        cls.invoices = {entry['id']: entry for entry in cls.data['invoices']}
        cls.paths = originals()

    def one_file(self, pattern):
        matches = list(CASE.glob(pattern))
        self.assertEqual(len(matches), 1, f'{pattern}: erwartet genau eine Originaldatei')
        return matches[0]

    def test_inventory_has_all_project_documents_and_no_export_debris(self):
        self.assertEqual(len(self.paths), 198)
        primary = [p for p in self.paths if p.suffix != '.xml']
        self.assertEqual(Counter(int(p.name[:3]) for p in primary), Counter({n: 1 for n in PRIMARY_NUMBERS}))
        self.assertEqual(sum(p.suffix == '.xml' for p in self.paths), 28)
        self.assertEqual({p.suffix for p in self.paths}, FORMATS)
        self.assertEqual(len({p.stem.casefold() for p in self.paths}), len(self.paths))
        for p in self.paths:
            self.assertTrue(p.name.isascii(), p.name)
            self.assertGreater(p.stat().st_size, 100, p.name)
        extras = {p.name for p in CASE.iterdir() if p.is_file()} - {p.name for p in self.paths}
        self.assertEqual(extras, {'README.md', 'rubric.yaml'})
        self.assertEqual({p.name for p in CASE.iterdir() if p.is_dir()}, {'gesamt-pdf'})

    def test_finance_totals_and_each_bank_movement_reconcile(self):
        d = self.data
        self.assertEqual(len(d['invoices']), 38)
        self.assertEqual(len(self.invoices), 38, 'Doppelte Forderungs-ID in kanonischen Daten')
        self.assertEqual(sum(amount(r['gross']) for r in d['invoices']), FINAL_COST)
        self.assertEqual(amount(d['total']), FINAL_COST)
        self.assertEqual(amount(d['base']), BASE_COST)
        self.assertEqual(sum(amount(g[2]) for g in d['groups']), BASE_COST)
        self.assertEqual(FINAL_COST - BASE_COST, Decimal('18445.00'))
        self.assertEqual(amount(d['equity']), Decimal('2100000.00'))
        self.assertEqual(amount(d['loan']), Decimal('1600000.00'))
        self.assertEqual(amount(d['equity']) + amount(d['loan']) - FINAL_COST, FINAL_CASH)

        bank = d['transactions']
        self.assertEqual(len(bank), 46)
        self.assertEqual(len({t['reference'] for t in bank}), len(bank), 'Bankreferenzen sind nicht eindeutig')
        self.assertEqual([t['date'] for t in bank], sorted(t['date'] for t in bank))
        running = Decimal(0)
        settled = defaultdict(Decimal)
        funding = Decimal(0)
        for t in bank:
            with self.subTest(reference=t['reference']):
                value = amount(t['amount'])
                self.assertGreater(value, 0)
                self.assertIn(t['direction'], {'Abfluss', 'Zufluss', 'Erstattung'})
                running += -value if t['direction'] == 'Abfluss' else value
                self.assertEqual(amount(t['balance']), running)
                if t['direction'] == 'Zufluss':
                    funding += value
                else:
                    self.assertIn(t['invoice'], self.invoices)
                    settled[t['invoice']] += value if t['direction'] == 'Abfluss' else -value
        self.assertEqual(running, FINAL_CASH)
        self.assertEqual(funding, Decimal('3700000.00'))
        self.assertEqual(set(settled), set(self.invoices))
        for ident, invoice in self.invoices.items():
            self.assertEqual(settled[ident], amount(invoice['gross']), ident)
            self.assertEqual(amount(invoice['net']) + amount(invoice['vat']), amount(invoice['gross']), ident)
            text = pdf_text(CASE/invoice['filename'])
            # Kaufpreis und Versicherungsbeitrag sind keine Umsatzsteuerrechnungen;
            # ihr interner Buchungsschlüssel verweist auf Urkunde bzw. Vertrag.
            source_reference = {'KP-2026-118': 'UVZ 118/2026', 'BV-27-180': 'BV-SW-2027-180'}.get(ident, ident)
            self.assertIn(source_reference, text)
            self.assertIn(german(invoice['gross']), text, invoice['filename'])
        # Tatsächliche Doppelzahlung und spätere Rückzahlung müssen sichtbar bleiben.
        electro = [t for t in bank if t['invoice'] == 'LE-28-061']
        self.assertEqual(Counter(t['direction'] for t in electro), {'Abfluss': 2, 'Erstattung': 1})
        self.assertEqual({amount(t['amount']) for t in electro}, {Decimal('89250.00')})

    def test_xml_and_pdf_are_two_representations_of_one_receivable(self):
        expected = {r['id']: r for r in self.data['invoices']
                    if r['date'] >= '2027-01-01' and (r['rate'] != 0 or r['id'] == 'LE-PV-28-099')}
        self.assertEqual(len(expected), 28)
        seen = set()
        xml_receivables = Decimal(0)
        for path in [p for p in self.paths if p.suffix == '.xml']:
            with self.subTest(file=path.name):
                xml = ET.parse(path).getroot()
                self.assertIn(xml.tag, {
                    '{urn:oasis:names:specification:ubl:schema:xsd:Invoice-2}Invoice',
                    '{urn:oasis:names:specification:ubl:schema:xsd:CreditNote-2}CreditNote'})
                ident = xml.findtext('cbc:ID', namespaces=NS)
                self.assertIn(ident, expected)
                self.assertNotIn(ident, seen, 'Doppelte strukturierte Forderung')
                seen.add(ident)
                inv = expected[ident]
                self.assertEqual(path.name, Path(inv['filename']).stem + '_XRechnung.xml')
                self.assertEqual(xml.findtext('cbc:IssueDate', namespaces=NS), inv['date'])
                self.assertEqual(xml.findtext('cbc:DocumentCurrencyCode', namespaces=NS), 'EUR')
                self.assertEqual(xml.findtext('cbc:BuyerReference', namespaces=NS), 'SW-HI-26-08')
                self.assertIn('en16931', xml.findtext('cbc:CustomizationID', default='', namespaces=NS))
                self.assertEqual(xml.findtext('cac:AccountingCustomerParty/cac:Party/cac:PartyLegalEntity/cbc:RegistrationName', namespaces=NS), 'Steinbogen Wohnen GmbH')

                credit = xml.tag.endswith('}CreditNote')
                self.assertEqual(credit, amount(inv['gross']) < 0)
                sign = Decimal(-1 if credit else 1)
                total = xml.find('cac:LegalMonetaryTotal', NS)
                self.assertIsNotNone(total)
                def value(field, default=None):
                    node = total.find('cbc:' + field, NS)
                    if node is None and default is not None:
                        return default
                    self.assertIsNotNone(node, field)
                    self.assertEqual(node.get('currencyID'), 'EUR')
                    return amount(node.text)
                net = value('TaxExclusiveAmount')
                gross = value('TaxInclusiveAmount')
                prepaid = value('PrepaidAmount', Decimal(0))
                payable = value('PayableAmount')
                tax = amount(xml.findtext('cac:TaxTotal/cbc:TaxAmount', namespaces=NS))
                prior = [self.invoices[i] for i in inv.get('prior') or []]
                self.assertEqual(gross - prepaid, payable, 'BT-115 muss Abschläge absetzen')
                self.assertEqual(gross - net, tax)
                self.assertEqual(sign * payable, amount(inv['gross']))
                self.assertEqual(prepaid, sum((amount(p['gross']) for p in prior), Decimal(0)))
                self.assertEqual(sign * net, amount(inv['net']) + sum((amount(p['net']) for p in prior), Decimal(0)))
                self.assertEqual(sign * tax, amount(inv['vat']) + sum((amount(p['vat']) for p in prior), Decimal(0)))
                line_tag = 'cac:CreditNoteLine' if credit else 'cac:InvoiceLine'
                lines = xml.findall(line_tag, NS)
                self.assertGreater(len(lines), 0)
                self.assertEqual(sum(amount(n.findtext('cbc:LineExtensionAmount', namespaces=NS)) for n in lines), value('LineExtensionAmount'))
                self.assertEqual(value('LineExtensionAmount'), net)
                for line in lines:
                    quantity = line.find('cbc:CreditedQuantity' if credit else 'cbc:InvoicedQuantity', NS)
                    self.assertIsNotNone(quantity)
                    price = Decimal(line.findtext('cac:Price/cbc:PriceAmount', namespaces=NS))
                    self.assertEqual(amount(Decimal(quantity.text) * price), amount(line.findtext('cbc:LineExtensionAmount', namespaces=NS)))
                pdf = pdf_text(CASE/inv['filename'])
                self.assertIn(ident, pdf)
                self.assertIn(german(sign * payable), pdf)
                xml_receivables += sign * payable
        self.assertEqual(seen, set(expected))
        other = sum((amount(r['gross']) for r in self.data['invoices'] if r['id'] not in seen), Decimal(0))
        self.assertEqual(xml_receivables + other, FINAL_COST, 'XML/PDF dürfen keine doppelten Kosten erzeugen')
        pv = self.invoices['LE-PV-28-099']
        self.assertEqual((amount(pv['net']), amount(pv['vat'])), (Decimal('45000.00'), Decimal('0.00')))
        self.assertEqual(amount(self.invoices['LE-28-061']['net']) + amount(self.invoices['LE-28-098']['net']), Decimal('105000.00'))

    def test_eight_rental_contracts_have_matching_notices_and_handovers(self):
        self.assertEqual(len(list(CASE.glob('*_Mietvertrag_WE*.docx'))), 8)
        self.assertEqual(len(list(CASE.glob('*_Uebergabe_WE*.pdf'))), 8)
        self.assertEqual(len(list(CASE.glob('*_Vorvertragliche_Auskunft_WE*.eml'))), 8)
        self.assertEqual(self.data['areas'], AREAS)
        self.assertEqual(sum(AREAS), 600)
        recipients = set()
        for unit, area in enumerate(AREAS, 1):
            with self.subTest(wohnung=unit):
                contract = docx_text(self.one_file(f'*_Mietvertrag_WE{unit:02d}.docx'))
                handover = pdf_text(self.one_file(f'*_Uebergabe_WE{unit:02d}.pdf'))
                notice = email_message(self.one_file(f'*_Vorvertragliche_Auskunft_WE{unit:02d}.eml'))
                text = clean(notice.get_body(preferencelist=('plain',)).get_content())
                self.assertIn(f'Wohnung {unit:02d}', contract)
                self.assertIn(f'Wohnung {unit:02d}', handover)
                self.assertIn(f'Wohnung {unit:02d}', text)
                self.assertRegex(contract, rf'\b{area}\s*m²')
                self.assertRegex(handover, rf'\b{area}\s*m²')
                rent = Decimal(area) * Decimal('13.50')
                self.assertIn(german(rent), contract)
                self.assertIn(german(rent), text)
                self.assertIn(german(rent * 3), contract)
                self.assertIn('drei gleichen monatlichen Raten', contract)
                self.assertIn('01.10.2028', contract)
                self.assertIn('29.09.2028', handover)
                self.assertIn('01.09.2028', contract)
                self.assertLess(parsedate_to_datetime(str(notice['Date'])).date(), date(2028, 9, 1))
                self.assertIn('01.10.2014', text)
                match = re.search(r'vermietet an (.+?) die Wohnung', contract)
                self.assertIsNotNone(match)
                self.assertIn(match.group(1), handover)
                self.assertEqual(notice['To'].addresses[0].display_name, match.group(1))
                recipients.add(notice['To'].addresses[0].addr_spec)
                for kind in 'EKWH':
                    self.assertIn(f'SW-{kind}-{unit:02d}', handover)
        self.assertEqual(len(recipients), 8)

    def test_eml_single_parsed_sender_recipient_and_unchanged_mime_attachments(self):
        attachment_count = 0
        message_ids = set()
        for path in [p for p in self.paths if p.suffix == '.eml']:
            with self.subTest(file=path.name):
                msg = email_message(path)
                self.assertEqual(msg.defects, [])
                for header in ('From', 'To'):
                    self.assertEqual(len(msg.get_all(header, [])), 1)
                    self.assertEqual(msg[header].defects, ())
                    self.assertEqual(len(msg[header].addresses), 1, f'{path.name}: {header} enthält mehrere Adressen')
                    address = msg[header].addresses[0]
                    self.assertTrue(address.username)
                    self.assertTrue(address.domain.endswith('.example'), address.addr_spec)
                for field in ('Date', 'Subject', 'Message-ID', 'Return-Path', 'Received', 'MIME-Version'):
                    self.assertTrue(msg[field], field)
                self.assertIsNotNone(parsedate_to_datetime(str(msg['Date'])).tzinfo)
                self.assertNotIn(str(msg['Message-ID']), message_ids)
                message_ids.add(str(msg['Message-ID']))
                plain = msg.get_body(preferencelist=('plain',))
                self.assertIsNotNone(plain)
                body = plain.get_content()
                self.assertGreater(len(body.strip()), 300)
                declared = re.findall(r'^Anlagen:\s*(.*)$', body, re.M)
                self.assertEqual(len(declared), 1, 'Genau ein Anlagenverzeichnis im Nachrichtentext')
                expected = [] if declared[0].strip() == 'keine' else [n.strip() for n in declared[0].split(';')]
                parts = list(msg.iter_attachments())
                self.assertEqual([part.get_filename() for part in parts], expected)
                for part in parts:
                    name = part.get_filename()
                    self.assertEqual(Path(name).name, name)
                    self.assertTrue((CASE/name).is_file(), name)
                    self.assertTrue(part.get_payload(decode=True) == (CASE/name).read_bytes(), f'{path.name}: veralteter MIME-Anhang {name}')
                    self.assertEqual(part.get_content_disposition(), 'attachment')
                    attachment_count += 1
        self.assertGreater(attachment_count, 0, 'Anlagen wurden nur erwähnt, aber nicht eingebettet')

    def test_native_formats_open_and_contain_content(self):
        for path in self.paths:
            with self.subTest(file=path.name):
                kind = path.suffix
                if kind == '.pdf':
                    self.assertTrue(path.read_bytes().startswith(b'%PDF-'))
                    reader = PdfReader(path)
                    self.assertFalse(reader.is_encrypted)
                    self.assertGreater(len(reader.pages), 0)
                    self.assertEqual(reader.metadata.author, 'Klotzkette')
                    text = pdf_text(path)
                    self.assertGreater(len(text), 100)
                    self.assertNotIn(NOTICE_DE, text)
                    self.assertNotIn(NOTICE_EN, text)
                elif kind in {'.docx', '.xlsx'}:
                    self.assertTrue(zipfile.is_zipfile(path))
                    with zipfile.ZipFile(path) as z:
                        self.assertIn('[Content_Types].xml', z.namelist())
                        core = ET.fromstring(z.read('docProps/core.xml'))
                        self.assertEqual(core.findtext('dc:creator', namespaces=NS), 'Klotzkette')
                        member = 'word/document.xml' if kind == '.docx' else 'xl/workbook.xml'
                        ET.fromstring(z.read(member))
                    if kind == '.docx':
                        self.assertGreater(len(docx_text(path)), 500)
                elif kind == '.png':
                    with Image.open(path) as image:
                        self.assertEqual(image.format, 'PNG')
                        self.assertGreaterEqual(min(image.size), 1000)
                        self.assertGreater(max(ImageStat.Stat(image.convert('RGB')).stddev), 10)
                        self.assertEqual(image.info.get('Author'), 'Klotzkette')
                elif kind == '.csv':
                    with path.open(encoding='utf-8-sig', newline='') as f:
                        rows = list(csv.reader(f))
                    self.assertGreater(len(rows), 2)
                    self.assertGreater(len(rows[0]), 1)
                    self.assertTrue(all(len(row) == len(rows[0]) for row in rows), path.name)
                elif kind == '.txt':
                    self.assertGreater(len(path.read_text(encoding='utf-8')), 200)
                elif kind == '.xml':
                    ET.parse(path)

    def test_workbooks_cache_real_formulas_and_count_each_invoice_once(self):
        books = {}
        try:
            for path in [p for p in self.paths if p.suffix == '.xlsx']:
                formulas = load_workbook(path, data_only=False)
                values = load_workbook(path, data_only=True)
                books[int(path.name[:3])] = values
                count = 0
                try:
                    for sheet in formulas:
                        for row in sheet:
                            for cell in row:
                                cached = values[sheet.title][cell.coordinate]
                                self.assertNotEqual(cached.data_type, 'e', f'{path.name}: {sheet.title}!{cell.coordinate}')
                                if cell.data_type == 'f':
                                    count += 1
                                    self.assertIsNotNone(cached.value, f'{path.name}: Formel ohne Cache {cell.coordinate}')
                    self.assertGreater(count, 20, path.name)
                finally:
                    formulas.close()
                if 'Belege' in values:
                    rows = [r for r in values['Belege'].iter_rows(values_only=True) if r[0] in self.invoices]
                    self.assertEqual(Counter(r[0] for r in rows), Counter({i: 1 for i in self.invoices}))
                    for row in rows:
                        inv = self.invoices[row[0]]
                        self.assertEqual(tuple(amount(v) for v in row[4:7]), tuple(amount(inv[k]) for k in ('net', 'vat', 'gross')))
                    self.assertEqual(sum(amount(r[6]) for r in rows), FINAL_COST)
            self.assertEqual(set(books), {120, 121, 122})
            self.assertEqual(amount(books[120]['Kostenstand']['D13'].value), FINAL_COST)
            self.assertEqual(amount(books[122]['Finanzierung']['B8'].value), FINAL_COST)
            self.assertEqual(amount(books[122]['Finanzierung']['B12'].value), FINAL_CASH)
            bank_rows = [r for r in books[121]['Bank'].iter_rows(values_only=True)
                         if r[1] in {t['reference'] for t in self.data['transactions']}]
            self.assertEqual(len(bank_rows), 46)
            self.assertEqual(amount(bank_rows[-1][5]), FINAL_CASH)
            self.assertEqual(amount(books[122]['Vermietung']['F14'].value), Decimal('102000.00'))
        finally:
            for book in books.values():
                book.close()

    def assert_export_contains_original(self, original, pages):
        """Prüft Inhalte, ohne plattformabhängige Office-Seitenumbrüche festzuschreiben."""
        self.assertGreater(len(pages), 0)
        rendered = '\n'.join(page.extract_text() or '' for page in pages)
        for notice in (NOTICE_DE, NOTICE_EN):
            self.assertNotIn(notice, clean(rendered))
        if original.suffix == '.pdf':
            source = PdfReader(original)
            self.assertEqual(len(pages), len(source.pages))
            for before, after in zip(source.pages, pages):
                source_size = sorted(float(v) for v in (before.mediabox.width, before.mediabox.height))
                is_a4 = all(abs(actual - expected) < 1 for actual, expected in zip(source_size, (595.276, 841.890)))
                if is_a4:
                    self.assertEqual(list(before.mediabox), list(after.mediabox))
                    self.assertEqual(before.get_contents().get_data(), after.get_contents().get_data(),
                                     'A4-Original-PDF-Inhalt beim Export verändert')
                else:
                    # Technische A3-Originale liegen unverändert im Akten-ZIP;
                    # die beiden PDF-Lesefassungen vereinheitlichen das Papier auf A4.
                    self.assertTrue(all(abs(actual - expected) < 1 for actual, expected in zip(source_size, (841.890, 1190.551))))
                    exported_size = sorted(float(v) for v in (after.mediabox.width, after.mediabox.height))
                    self.assertTrue(all(abs(actual - expected) < 1 for actual, expected in zip(exported_size, (595.276, 841.890))))
                    self.assertEqual(clean(before.extract_text() or ''), clean(after.extract_text() or ''),
                                     'A3-Planinhalt bei der A4-Lesefassung verloren')
            return
        if original.suffix == '.png':
            with Image.open(original) as source:
                source = source.convert('RGB')
                images = [image.image.convert('RGB') for page in pages for image in page.images]
                self.assertTrue(any(image.size == source.size and image.tobytes() == source.tobytes()
                                    for image in images), 'Originalbild fehlt oder wurde im Export verändert')
            return
        if original.suffix == '.docx':
            expected = docx_text(original)
        elif original.suffix == '.eml':
            message = email_message(original)
            expected = (str(message['Subject']) + ' ' + message['From'].addresses[0].addr_spec + ' '
                        + message['To'].addresses[0].addr_spec + ' '
                        + message.get_body(preferencelist=('plain',)).get_content())
        elif original.suffix == '.xml':
            expected = ' '.join(ET.parse(original).getroot().itertext())
        elif original.suffix == '.csv':
            with original.open(encoding='utf-8-sig', newline='') as source:
                rows = list(csv.reader(source))
            # Schmale Tabellen umbrechen auch innerhalb eines Wortes oder Datums.
            # Tabellen und lange Datensatzansichten dürfen anders angeordnet sein;
            # jede Datenzelle muss dennoch vollständig in Originalreihenfolge stehen.
            def compact(value):
                return ''.join(re.findall(r'[^\W_]+', unicodedata.normalize('NFKC', value).casefold()))
            compact_rendered = compact(rendered)
            for heading in rows[0]:
                self.assertIn(compact(heading), compact_rendered)
            offset = 0
            for row in rows[1:]:
                for cell in row:
                    cell_text = compact(cell)
                    found = compact_rendered.find(cell_text, offset)
                    self.assertGreaterEqual(found, offset, f'CSV-Zelle fehlt im PDF: {cell}; Zeile: {row}')
                    offset = found + len(cell_text)
            return
        elif original.suffix == '.xlsx':
            book = load_workbook(original, data_only=True)
            try:
                # Beschriftungen, Beleg- und Bankreferenzen müssen aus allen Blättern
                # lesbar bleiben. Zahlen und Formeln prüft der eigenständige Finanztest.
                expected = ' '.join(cell.value for sheet in book for row in sheet for cell in row
                                    if isinstance(cell.value, str))
            finally:
                book.close()
        else:
            expected = original.read_text(encoding='utf-8')
        expected_words = words(expected)
        actual_words = words(rendered)
        retained = sum((expected_words & actual_words).values())
        required = sum(expected_words.values())
        self.assertGreater(required, 20)
        self.assertGreaterEqual(retained / required, 0.98,
                                f'{original.name}: Originalinhalt fehlt im PDF ({retained}/{required} Wörter)')

    def test_release_archives_exact_originals_flat_readme_first_and_pdf_coverage(self):
        if ASSETS is None:
            self.skipTest('Release-Pakete nur mit --assets DIR prüfen')
        aggregate = CASE/'gesamt-pdf'/f'{SLUG}_gesamt.pdf'
        self.assertTrue(aggregate.is_file())
        expected_originals = {p.name: p for p in self.paths}
        expected_originals[aggregate.name] = aggregate
        original_zip = ASSETS/f'testakte-{SLUG}.zip'
        single_zip = ASSETS/f'testakte-{SLUG}-einzelpdfs.zip'
        for archive in (original_zip, single_zip):
            self.assertTrue(archive.is_file(), str(archive))
            with zipfile.ZipFile(archive) as z:
                names = z.namelist()
                self.assertEqual(names[0], 'README.txt')
                self.assertEqual(len(names), len(set(n.casefold() for n in names)))
                self.assertTrue(all('/' not in n and '\\' not in n and not n.startswith('.') for n in names))
                notice = z.read('README.txt')
                self.assertTrue(notice.startswith(NOTICE_BYTES))
                notice.decode('utf-8')
                self.assertIsNone(z.testzip())
        with zipfile.ZipFile(original_zip) as z:
            self.assertEqual(set(z.namelist()), {'README.txt', *expected_originals})
            for name, path in expected_originals.items():
                self.assertTrue(z.read(name) == path.read_bytes(), f'Archiv enthält veraltete oder veränderte Bytes: {name}')

        total = PdfReader(aggregate)
        outline = [item for item in total.outline if not isinstance(item, list)]
        self.assertEqual([item.title for item in outline], ['Dokumentenregister'] + [p.name for p in self.paths])
        starts = [total.get_destination_page_number(item) for item in outline]
        self.assertEqual(starts[0], 0)
        self.assertGreater(starts[1], 0)
        self.assertEqual(starts, sorted(set(starts)))
        with zipfile.ZipFile(single_zip) as z:
            self.assertEqual(set(z.namelist()), {'README.txt', *(p.stem + '.pdf' for p in self.paths)})
            for i, original in enumerate(self.paths, 1):
                with self.subTest(rendered=original.name):
                    single = PdfReader(io.BytesIO(z.read(original.stem + '.pdf')))
                    end = starts[i + 1] if i + 1 < len(starts) else len(total.pages)
                    self.assertFalse(single.is_encrypted)
                    self.assert_export_contains_original(original, list(single.pages))
                    self.assert_export_contains_original(original, list(total.pages[starts[i]:end]))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--assets', type=Path, help='Verzeichnis der fertig gebauten ZIPs, zum Beispiel dist')
    args, remaining = parser.parse_known_args()
    ASSETS = args.assets.resolve() if args.assets else None
    unittest.main(argv=[sys.argv[0], *remaining])
