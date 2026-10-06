#!/usr/bin/env python3
"""Native Neuberechnung und positive/negative Eingabeprüfungen der drei WH26-Mappen."""
from pathlib import Path
import json,zipfile,subprocess,shutil,os,sys,tempfile,xml.etree.ElementTree as ET
from datetime import datetime
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
CASE=ROOT/'testakten/bauwirtschaft-baumanagement-werkhalle-warendorf'
QA=Path('/tmp/bauwirtschaft-vertiefung-20261006/warendorf')
SOFFICE=Path('/Users/klotzkette/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'; Q='{'+NS+'}'
FILES=['41_Energie_und_Freigaben.xlsx','42_Mobilversorgung_Kosten.xlsx','43_Zahlungen_und_Entscheidungen.xlsx']
AREAS=[['D28','D17'],['D26','D25'],['E26','H28']]

def edit(src,dest,changes=(),areas=None):
    with zipfile.ZipFile(src) as zin,zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            data=zin.read(i)
            if i.filename.startswith('xl/worksheets/sheet') and i.filename.endswith('.xml'):
                index=int(i.filename.split('sheet')[-1].split('.')[0]); s=ET.fromstring(data)
                for sheet,ref,value in changes:
                    if index!=sheet:continue
                    sd=s.find(Q+'sheetData');rn=int(''.join(c for c in ref if c.isdigit()))
                    row=next((r for r in sd if r.get('r')==str(rn)),None)
                    if row is None:row=ET.SubElement(sd,Q+'row',{'r':str(rn)})
                    c=next((c for c in row if c.get('r')==ref),None)
                    if c is None:c=ET.SubElement(row,Q+'c',{'r':ref})
                    for ch in list(c):c.remove(ch)
                    c.attrib.pop('t',None)
                    if isinstance(value,str):
                        c.set('t','inlineStr');ET.SubElement(ET.SubElement(c,Q+'is'),Q+'t').text=value
                    elif value is not None:ET.SubElement(c,Q+'v').text=str(value)
                for c in s.findall('.//'+Q+'c'):
                    if c.find(Q+'f') is not None:
                        for cached in c.findall(Q+'v'):c.remove(cached)
                if areas:
                    pr=s.find(Q+'sheetPr')
                    if pr is None:pr=ET.Element(Q+'sheetPr');s.insert(0,pr)
                    ps=pr.find(Q+'pageSetUpPr')
                    if ps is None:ps=ET.SubElement(pr,Q+'pageSetUpPr')
                    ps.set('fitToPage','1')
                    for tag in ['pageMargins','pageSetup']:
                        for element in s.findall(Q+tag):s.remove(element)
                    ET.SubElement(s,Q+'pageMargins',dict(left='0.3',right='0.3',top='0.3',bottom='0.3',header='0.1',footer='0.1'))
                    ET.SubElement(s,Q+'pageSetup',dict(paperSize='9',orientation='landscape',fitToWidth='1',fitToHeight='0'))
                data=ET.tostring(s,encoding='utf-8',xml_declaration=True)
            elif i.filename=='xl/workbook.xml':
                w=ET.fromstring(data)
                if areas:
                    dn=w.find(Q+'definedNames')
                    if dn is None:dn=ET.SubElement(w,Q+'definedNames')
                    for x in list(dn):
                        if x.get('name') in ['_xlnm.Print_Area','_xlnm.Print_Titles']:dn.remove(x)
                    for index,(s,end) in enumerate(zip(w.find(Q+'sheets'),areas)):
                        name=s.get('name').replace("'","''")
                        ET.SubElement(dn,Q+'definedName',dict(name='_xlnm.Print_Area',localSheetId=str(index))).text=f"'{name}'!$A$1:${end[0]}${end[1:]}"
                        ET.SubElement(dn,Q+'definedName',dict(name='_xlnm.Print_Titles',localSheetId=str(index))).text=f"'{name}'!$1:$5"
                calc=w.find(Q+'calcPr')
                if calc is None:calc=ET.SubElement(w,Q+'calcPr')
                calc.set('calcMode','auto');calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1')
                data=ET.tostring(w,encoding='utf-8',xml_declaration=True)
            elif i.filename=='docProps/core.xml':
                core=ET.fromstring(data)
                for tag,val in [('creator','Klotzkette'),('language','de-DE')]:
                    q='{http://purl.org/dc/elements/1.1/}'+tag
                    el=core.find(q)
                    if el is None:el=ET.SubElement(core,q)
                    el.text=val
                data=ET.tostring(core,encoding='utf-8',xml_declaration=True)
            zout.writestr(i,data)

def recalc(files,out,profile):
    out.mkdir(parents=True,exist_ok=True)
    cmd=[str(SOFFICE),'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out),*[str(p) for p in files]]
    p=subprocess.run(cmd,text=True,capture_output=True,timeout=180)
    (QA/(out.name+'.log')).write_text(p.stdout+p.stderr)
    assert all((out/f.name).exists() for f in files),p.stdout+p.stderr

def final_metadata(src,dest):
    """Office kann den Autor beim Speichern ersetzen; finale Metadaten danach setzen.

    Nur core.xml wird geändert. Insbesondere bleiben Formeln, native Ergebniswerte,
    Layout und Druckbereiche bytegleich zum geprüften Office-Ergebnis erhalten.
    """
    with zipfile.ZipFile(src) as zin,zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as zout:
        assert 'docProps/core.xml' in zin.namelist()
        for i in zin.infolist():
            data=zin.read(i)
            if i.filename=='docProps/core.xml':
                core=ET.fromstring(data)
                for tag,val in [('creator','Klotzkette'),('language','de-DE')]:
                    q='{http://purl.org/dc/elements/1.1/}'+tag
                    el=core.find(q)
                    if el is None:el=ET.SubElement(core,q)
                    el.text=val
                data=ET.tostring(core,encoding='utf-8',xml_declaration=True)
            zout.writestr(i,data)
    props=load_workbook(dest,read_only=True).properties
    assert (props.creator,props.language)==('Klotzkette','de-DE'),dest
    with zipfile.ZipFile(src) as zin,zipfile.ZipFile(dest) as zout:
        assert zin.namelist()==zout.namelist()
        assert all(zin.read(n)==zout.read(n) for n in zin.namelist() if n!='docProps/core.xml')

def value(file,sheet,cell):return load_workbook(file,data_only=True).worksheets[sheet-1][cell].value

def serial(s):return (datetime.fromisoformat(s)-datetime(1899,12,30)).days

def main():
    QA.mkdir(parents=True,exist_ok=True)
    result=[]
    cases=[
    ('energie_spaeter',0,[(1,'B6',serial('2026-11-23'))],[(1,'B16',datetime(2026,12,14))]),
    ('energie_freigabe_teil',0,[(2,'B6',1)],[(1,'B21','Freigaben offen')]),
    ('energie_freigabe_voll',0,[(2,f'B{r}',1) for r in range(6,10)],[(1,'B21','Freigaben dokumentiert')]),
    ('energie_abend_offen',0,[(1,'B11',2)]+[(2,f'B{r}',1) for r in range(6,10)],[(1,'B14',datetime(2026,10,30)),(1,'B21','Freigaben offen')]),
    ('energie_fehlendes_datum',0,[(1,'B7',None)],[(1,'B14','Eingabe offen')]),
    ('energie_negative_dauer',0,[(1,'B8',-1)],[(1,'B14','Eingabe offen')]),
    ('energie_inaktive_dauer_leer',0,[(1,'B10',None)],[(1,'B14',datetime(2026,11,2))]),
    *[(name,0,[(1,'B11',invalid)],[(1,'B12','Variante prüfen'),(1,'B13','Variante prüfen'),(1,'B14','Eingabe offen'),(1,'B16','Eingabe offen'),(1,'B21','Variante prüfen')]) for name,invalid in [('energie_variante_leer',None),('energie_variante_text','unbekannt'),('energie_variante_ausserhalb',3),('energie_variante_dezimal',1.5)]],
    ('kosten_abend_unbekannt',1,[(1,'B6',2)],[(1,'B9',7730),(1,'B12','Preis offen')]),
    ('kosten_abend_bepreist',1,[(1,'B6',2),(2,'B19',2000)],[(1,'B12',45730),(1,'B14',54418.7)]),
    ('kosten_abend_null',1,[(1,'B6',2),(2,'B19',0)],[(1,'B12',43730)]),
    ('kosten_verlaengerung',1,[(2,'B16',2)],[(1,'B12',45000),(1,'B17',40698)]),
    ('kosten_verlaengerung_zu_lang',1,[(2,'B16',3)],[(1,'B12','Preis offen')]),
    ('kosten_fehlender_satz',1,[(1,'B6',2),(2,'B19',0),(2,'B12',None)],[(1,'B12','Preis offen')]),
    ('zahlung_reserve_september',2,[(2,'C6',serial('2026-09-28'))],[(1,'E6',322000),(1,'E9',124860)]),
    ('zahlung_mobil_einplanen',2,[(2,'D10',1),(2,'D11',1)],[(1,'E9',82020),(1,'B17',0)]),
    ('zahlung_einzeln_freigegeben',2,[(2,'E6',1)],[(1,'B17',190400)]),
    ('zahlung_fehlendes_datum',2,[(2,'D12',1)],[(1,'E9','Eingabe offen'),(2,'H12','Datum fehlt')]),
    ('zahlung_fehlender_preis',2,[(2,'D15',1),(2,'C15',serial('2026-11-30'))],[(1,'E9','Eingabe offen'),(2,'H15','Betrag fehlt')]),
    ('zahlung_doppelter_schluessel',2,[(2,'A7','WH 260924')],[(1,'E9','Eingabe offen'),(2,'H6','ID doppelt')]),
    ('zahlung_neuer_vorgang',2,[(2,'A16','WH-Z16'),(2,'B16',1000),(2,'C16',serial('2026-12-15')),(2,'D16',1),(2,'E16',0),(2,'F16',0)],[(1,'E9',123670)]),
    ('zahlung_zeitgrenze',2,[(2,'C9',serial('2027-01-01'))],[(1,'E9','Eingabe offen'),(2,'H9','Außerhalb Zeitraum')]),
    ('zahlung_negativer_betrag',2,[(2,'B6',-100)],[(1,'E9','Eingabe offen')]),
    ('zahlung_abend_ohne_service',2,[(2,'D13',1)],[(1,'E9','Eingabe offen')]),
    ('zahlung_invalides_planfeld',2,[(2,'D6',3)],[(1,'E9','Eingabe offen')]),
    ('zahlung_fehlende_steuer',2,[(1,'B13',None)],[(1,'E9','Eingabe offen')]),
    ('zahlung_druckluft_fortschreiben',2,[(1,'B20',32000),(2,'D14',1)],[(2,'B14',2000),(1,'E9',122480)]),
    ]
    with tempfile.TemporaryDirectory(prefix='wh26-native-',dir='/tmp') as tmp:
        temp=Path(tmp);src=temp/'original';src.mkdir()
        for i,n in enumerate(FILES):edit(CASE/n,src/n,areas=AREAS[i])
        recalc(list(src.glob('*.xlsx')),temp/'baseline',temp/'profile-base')
        base=temp/'baseline'
        for n in FILES:
            b=load_workbook(base/n,data_only=True)
            errors=[(s.title,c.coordinate,c.value) for s in b for row in s for c in row if c.data_type=='e']
            assert not errors,errors
        assert value(base/FILES[0],1,'B14')==datetime(2026,11,2)
        assert value(base/FILES[0],1,'B16')==datetime(2026,12,7)
        assert value(base/FILES[1],1,'B12')==36000
        assert value(base/FILES[1],1,'B16')==12852
        assert value(base/FILES[2],1,'E9')==124860
        edits=temp/'mutations';edits.mkdir()
        for name,i,changes,expected in cases:edit(base/FILES[i],edits/(name+'.xlsx'),changes)
        recalc(list(edits.glob('*.xlsx')),temp/'results',temp/'profile-mutations')
        for name,i,changes,expected in cases:
            for sn,c,want in expected:
                got=value(temp/'results'/(name+'.xlsx'),sn,c)
                assert got==want or isinstance(want,float) and abs(got-want)<0.00001,(name,sn,c,got,want)
            result.append({'Fall':name,'Ergebnis':'bestanden','Prüfzellen':len(expected)})
        for n in FILES:final_metadata(base/n,CASE/n)
    (QA/'native_mutationstests.json').write_text(json.dumps({'Engine':str(SOFFICE),'Baseline':'bestanden','Metadaten_nach_Office':{'creator':'Klotzkette','language':'de-DE','Dateien':FILES,'Nur_core_xml_geaendert':True},'Mutationen':result},ensure_ascii=False,indent=2))
    print(f'Native Neuberechnung bestanden; {len(cases)} positive und negative Mutationen bestanden.')

if __name__=='__main__':main()
