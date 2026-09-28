#!/usr/bin/env python3
"""Build only the eleven native source documents of the Hessen energy case."""

from __future__ import annotations

import argparse
import csv
import json
from datetime import datetime, timezone
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from pathlib import Path
from xml.sax.saxutils import escape


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "scripts/fixtures/vergesellschaftung-artikel-15/case.json"
CASE = ROOT / "testakten/vergesellschaftung-energienetz-hessen"


def load_case():
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    names = [entry["file"] for entry in data["documents"]]
    if len(names) != 11 or len(set(names)) != 11:
        raise ValueError("Expected eleven distinct case documents")
    for entry in data["documents"]:
        name = Path(entry["file"])
        if name.name != entry["file"] or name.suffix != "." + entry["format"]:
            raise ValueError("Invalid case filename")
    if chr(167) in json.dumps(data, ensure_ascii=False):
        raise ValueError("Spell out Paragraf in the authored source")
    return data


def email_bytes(entry):
    message = EmailMessage(policy=SMTP)
    for key in ("from", "to", "cc", "subject"):
        if key in entry:
            message[key] = entry[key]
    message["Date"] = format_datetime(datetime.fromisoformat(entry["date"]))
    message["Message-ID"] = "<" + entry["message_id"] + ">"
    message.set_content("\n\n".join(entry["paragraphs"]) + "\n", charset="utf-8")
    return message.as_bytes()


def make_docx(entry, target):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    document.settings.odd_and_even_pages_header_footer = False
    section = document.sections[0]
    section.different_first_page_header_footer = False
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.1)
    section.left_margin, section.right_margin = Cm(2.3), Cm(2.3)
    section.header_distance, section.footer_distance = Cm(1.1), Cm(1.0)
    for name in ("Normal", "Title", "Heading 1", "Header", "Footer"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11 if name not in ("Header", "Footer") else 9)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + script), "Times New Roman")
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.13
        style.paragraph_format.widow_control = True
    for name in ("Title", "Heading 1"):
        document.styles[name].font.bold = True
        document.styles[name].paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = entry["sender"]
    footer = section.footer.paragraphs[0]
    footer.add_run(entry["reference"].split("|")[0].strip() + " | Seite ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = entry["title"]
    document.core_properties.author = entry["sender"].split("|")[0].strip()
    document.core_properties.last_modified_by = document.core_properties.author
    stamp = datetime.strptime(entry["date"], "%d.%m.%Y").replace(tzinfo=timezone.utc)
    document.core_properties.created = stamp
    document.core_properties.modified = stamp
    for index, page in enumerate(entry["pages"]):
        if index:
            document.add_page_break()
        else:
            document.add_paragraph(entry["title"], "Title")
            document.add_paragraph("An: " + entry["recipient"])
            document.add_paragraph(entry["date"] + " | " + entry["reference"])
        document.add_paragraph(page["heading"], "Heading 1")
        for text in page["paragraphs"]:
            document.add_paragraph(text)
    document.save(target)


def make_pdf(entry, target, font_dir):
    from akten_build_runtime import serif_font_path
    from pypdf import PdfReader
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    for bold, (name, filename) in enumerate((("CaseTimes", "Times New Roman.ttf"),
                                            ("CaseTimesBold", "Times New Roman Bold.ttf"))):
        font = font_dir / filename if font_dir else serif_font_path(bool(bold))
        if not font.is_file():
            raise FileNotFoundError(f"Provide Times New Roman fonts with --font-dir: {font}")
        pdfmetrics.registerFont(TTFont(name, str(font)))
    body = ParagraphStyle("Body", fontName="CaseTimes", fontSize=11, leading=13.3,
                          spaceAfter=8, allowWidows=0, allowOrphans=0)
    heading = ParagraphStyle("Heading", parent=body, fontName="CaseTimesBold",
                             spaceBefore=5, spaceAfter=10, keepWithNext=True)
    title = ParagraphStyle("Title", parent=heading, spaceBefore=0)
    story = []
    for index, page in enumerate(entry["pages"]):
        if index:
            story.append(PageBreak())
        else:
            story.extend([
                Paragraph(escape(entry["title"]), title),
                Paragraph(escape("An: " + entry["recipient"]), body),
                Paragraph(escape(entry["date"] + " | " + entry["reference"]), body),
                Spacer(1, 3 * mm),
            ])
        story.append(Paragraph(escape(page["heading"]), heading))
        story.extend(Paragraph(escape(text).replace("\n", "<br/>"), body)
                     for text in page["paragraphs"])

    def frame(canvas, document):
        canvas.saveState()
        canvas.setFont("CaseTimes", 9)
        canvas.drawString(23 * mm, A4[1] - 13 * mm, entry["sender"])
        canvas.drawString(23 * mm, 12 * mm, entry["reference"].split("|")[0].strip())
        canvas.drawRightString(A4[0] - 23 * mm, 12 * mm, f"Seite {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(str(target), pagesize=A4, leftMargin=23 * mm,
                                 rightMargin=23 * mm, topMargin=22 * mm,
                                 bottomMargin=21 * mm, title=entry["title"],
                                 author=entry["sender"].split("|")[0].strip())
    document.build(story, onFirstPage=frame, onLaterPages=frame)
    actual = len(PdfReader(target).pages)
    if actual != len(entry["pages"]):
        raise ValueError(f"Unexpected pagination in {target.name}: {actual}")


def build(output, font_dir):
    data = load_case()
    output.mkdir(parents=True, exist_ok=True)
    for entry in data["documents"]:
        target = output / entry["file"]
        kind = entry["format"]
        if kind == "eml":
            target.write_bytes(email_bytes(entry))
        elif kind == "txt":
            target.write_text(entry["title"] + "\n\n" + "\n\n".join(entry["paragraphs"])
                              + "\n", encoding="utf-8")
        elif kind == "csv":
            # These are raw source exports, not calculated Excel workbooks.
            with target.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream, delimiter=";", lineterminator="\r\n")
                writer.writerow(entry["columns"])
                writer.writerows(entry["rows"])
        elif kind == "docx":
            make_docx(entry, target)
        elif kind == "pdf":
            make_pdf(entry, target, font_dir)
        else:
            raise ValueError(kind)
        print(target.name)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=CASE)
    parser.add_argument("--font-dir", type=Path, help="Optionaler Ordner mit Times New Roman; sonst portable Schriftwahl")
    args = parser.parse_args()
    build(args.output.resolve(), args.font_dir.resolve() if args.font_dir else None)


if __name__ == "__main__":
    main()
