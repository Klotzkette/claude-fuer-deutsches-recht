#!/usr/bin/env python3
"""Sichert Mitarbeiterwege und die drei Fassungen der sechs Notariatsvorgänge."""

import importlib.util
import io
from pathlib import Path
import re
import tempfile
import unittest
import zipfile

from docx import Document
from pypdf import PdfReader
from testakte_disclaimer import NOTICE_BYTES, pdf_content_errors
from testakte_einzelpdf_common import expected_arcnames

ROOT = Path(__file__).resolve().parent.parent
PLUGIN = ROOT / 'notariat-alltag'
CASES = ["notariat-bautraegerkauf-parkhof-potsdam","notariat-grundschuld-bankauftrag-erfurt","notariat-unterschriftsbeglaubigung-hannover","notariat-gmbh-gruendung-rostock","notariat-kapitalerhoehung-geschaeftsfuehrer-ulm","notariat-anteilsuebertragung-verpfaendung-bremen"]
CORE = ["kaltstart-triage","notariat-002-beurkundung-ubeglaubigung-richtig-einordnen","notariat-023-identitaetspruefung-videoident-praesenztermin","gmbh-gruendung-gesellschafterliste","notariat-032-kapitalerhoehung-bar-sache-bezugsrecht","notariat-006-hr-anmeldung-gf-bestellung-abberufung-vertretung","gmbh-anteile-uebertragen-verpfaenden","bautraegervertrag-mabv-familiengesellschaft","grundschuld-buchgrundschuld-treuhand","qualitaetsgate-signatur-notarielle"]


def builder(filename):
    spec = importlib.util.spec_from_file_location(filename.replace('-', '_'), ROOT / 'scripts' / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class NotariatswerkstattTests(unittest.TestCase):
    def test_manifest_prose_displays_umlauts_without_changing_the_manifest(self):
        import json
        manifest = json.loads((PLUGIN / '.claude-plugin/plugin.json').read_text(encoding='utf-8'))
        display = builder('inject-direkt-loslegen-section.py').markdown_text(manifest['description'])
        for word in ('GmbH-Gründung', 'Kapitalerhöhung', 'Geschäftsführerwechsel', 'Anteilsgeschäfte'):
            self.assertIn(word, display)
        self.assertIn('GmbH-Gruendung', manifest['description'])

    def test_core_routes_are_available_and_use_decimal_workflows(self):
        self.assertEqual(len(CORE), 10)
        readme = (PLUGIN / 'README.md').read_text(encoding='utf-8')
        for slug in CORE:
            with self.subTest(skill=slug):
                text = (PLUGIN / 'skills' / slug / 'SKILL.md').read_text(encoding='utf-8')
                self.assertIn(slug, readme)
                self.assertIn('## 5. Ausgabeformat', text)
                self.assertIn('Times New Roman 11 pt', text)
                self.assertRegex(text, r'### 3[.]1[.] [^\n]+\n\n')
                self.assertNotRegex(text, r'(?m)^[0-9]+[.] [^\n]+$')
                self.assertNotIn('§', text)
                for target in re.findall(r'\]\(([^)]+)\)', text):
                    if '://' in target or target.startswith('#'):
                        continue
                    resolved = (PLUGIN / 'skills' / slug / target.partition('#')[0]).resolve()
                    self.assertTrue(resolved.is_relative_to(PLUGIN.resolve()), target)
                    self.assertTrue(resolved.is_file(), target)

    def test_prompts_preserve_form_distinctions_and_staff_boundary(self):
        for kind in ('werkstatt', 'schnellstart'):
            text = (PLUGIN / f'notariat-alltag-{kind}.md').read_text(encoding='utf-8')
            with self.subTest(prompt=kind):
                for anchor in ('Notariatsmitarbeiter', 'GmbHG Paragraf 53', 'Absatz 3', 'Paragraf 55 Absatz 1', 'Paragraf 1274', 'HGB Paragraf 12', 'MaBV Paragraf 3', 'Paragraf 794'):
                    self.assertIn(anchor, text)
                self.assertNotIn('skills/', text)
                self.assertIn('notariellen Prüfung', text)
                if kind == 'schnellstart':
                    self.assertLess(len(text.encode('utf-8')), 7500)
        legacy = (PLUGIN / 'skills/002-beurkundung-oder-unterschriftsbeglaubigung-richtig/SKILL.md').read_text(encoding='utf-8')
        self.assertNotIn('403 ZPO', legacy)
        self.assertNotIn('Einigungserklärung: Beurkundung nötig', legacy)
        entry = (PLUGIN / 'skills/kaltstart-triage/SKILL.md').read_text(encoding='utf-8')
        self.assertIn('nur den einen', entry)

    def test_every_case_has_eight_separate_native_documents(self):
        for slug in CASES:
            case = ROOT / 'testakten' / slug
            sources = [p for p in case.iterdir() if p.is_file() and p.name[:2].isdigit()]
            with self.subTest(case=slug):
                self.assertEqual(len(sources), 8)
                self.assertEqual({p.suffix for p in sources}, {'.docx', '.eml', '.txt'})
                self.assertEqual(sum(p.suffix == '.docx' for p in sources), 2)
                self.assertEqual(sum(p.suffix == '.eml' for p in sources), 4)
                for path in sources:
                    if path.suffix == '.docx':
                        text = '\n'.join(p.text for p in Document(path).paragraphs)
                        self.assertGreater(len(text), 1200, path.name)
                    else:
                        text = path.read_text(encoding='utf-8')
                    self.assertNotRegex(text.lower(), r'testakte|musterlösung|lösungsschlüssel|platzhalter|lorem ipsum')
                    self.assertNotIn('§', text)

    def test_original_date_survives_word_conversion(self):
        case = ROOT / 'testakten/notariat-kapitalerhoehung-geschaeftsfuehrer-ulm'
        doc = Document(case / '02_satzungsauszug.docx')
        self.assertEqual(doc.paragraphs[-1].text, '2. September 2026, 16:10 Uhr')
        self.assertFalse(doc.paragraphs[-1]._p.xpath('./w:pPr/w:numPr'))

    def test_combined_pdf_contains_every_source(self):
        for slug in CASES:
            case = ROOT / 'testakten' / slug
            path = case / 'gesamt-pdf' / f'{slug}_gesamt.pdf'
            with self.subTest(case=slug):
                self.assertEqual(pdf_content_errors(path.read_bytes()), [])
                text = '\n'.join(p.extract_text() or '' for p in PdfReader(path).pages)
                for source in case.iterdir():
                    if source.is_file() and source.name[:2].isdigit():
                        self.assertIn(source.name, text)

    def test_archives_are_flat_and_match_the_sources(self):
        for slug in CASES:
            case = ROOT / 'testakten' / slug
            sources = {p.name: p for p in case.iterdir() if p.is_file() and p.name[:2].isdigit()}
            expected = set(expected_arcnames(case))
            self.assertEqual(len(expected), 8)
            for filename in ('build-testakten-release-zips.py', 'build-testakten-einzelpdf-zips.py'):
                with self.subTest(case=slug, builder=filename), tempfile.TemporaryDirectory() as directory:
                    path, _ = builder(filename).build_single(case, Path(directory))
                    with zipfile.ZipFile(path) as archive:
                        names = archive.namelist()
                        self.assertEqual(len(names), len(set(names)))
                        self.assertTrue(all('/' not in n and chr(92) not in n for n in names))
                        self.assertFalse(any(n.endswith('.md') for n in names))
                        self.assertEqual(archive.read('README.txt'), NOTICE_BYTES)
                        self.assertIsNone(archive.testzip())
                        if 'einzelpdf' in filename:
                            self.assertEqual(set(names), expected | {'README.txt'})
                            for name in expected:
                                data = archive.read(name)
                                self.assertEqual(pdf_content_errors(data), [], name)
                                self.assertGreater(len(PdfReader(io.BytesIO(data)).pages), 0)
                        else:
                            self.assertEqual(set(names), set(sources) | {'README.txt', f'{slug}_gesamt.pdf'})
                            for name, source in sources.items():
                                self.assertEqual(archive.read(name), source.read_bytes())
                            self.assertEqual(archive.read(f'{slug}_gesamt.pdf'), (case / 'gesamt-pdf' / f'{slug}_gesamt.pdf').read_bytes())

    def test_cases_are_listed_in_plugin_and_central_index(self):
        for path in (PLUGIN / 'README.md', ROOT / 'testakten/README.md'):
            text = path.read_text(encoding='utf-8')
            for slug in CASES:
                self.assertIn(slug, text, path)


if __name__ == '__main__':
    unittest.main()
