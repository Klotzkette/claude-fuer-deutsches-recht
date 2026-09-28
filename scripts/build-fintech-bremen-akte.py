#!/usr/bin/env python3
"""Erzeugt die abgegrenzten Originalunterlagen des Bremer Darlehensverfahrens."""

from __future__ import annotations

import csv
import json
import re
from datetime import date, datetime
from email import policy
from email.message import EmailMessage
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, Table, TableStyle

from akten_build_runtime import serif_font_path
from testakte_disclaimer import NOTICE_MARKDOWN


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'scripts/fixtures/fintech-bremen'
CASE = json.loads((FIXTURES / 'case.json').read_text())
OUT = ROOT / 'testakten' / CASE['slug']
FONT = 'FintechSerif'
WIDTH = A4[0] - 104
PAYMENTS = [
    '2022-11-15', '2022-12-15', '2023-01-16', '2023-02-15',
    '2023-03-15', '2023-04-17', '2023-05-15', '2023-06-15',
    '2023-07-17', '2023-08-16', '2023-09-15', '2023-10-16',
    '2023-11-15', '2023-12-15', '2024-01-15', '2024-02-15',
    '2024-03-15', '2024-04-15',
]


def money(value):
    return f'{value:,.2f}'.replace(',', '_').replace('.', ',').replace('_', '.')


def table_data(page):
    data = page.get('table')
    if not data:
        return None
    return [data['headers'], *data['rows']]


def address(author):
    if author == 'Weserkontor Bank eG':
        return 'Langenstraße 18, 28195 Bremen | Geschäftskundenservice'
    for key in ('administrator', 'defence_firm', 'bank', 'fintech', 'debtor'):
        item = CASE[key]
        if any(item.get(field, '\0') in author for field in ('name', 'firm', 'lawyer')):
            return item['address']
    if 'Insolvenzgericht' in author:
        return 'Ostertorstraße 25, 28195 Bremen'
    if 'Amtsgericht' in author or 'Landgericht' in author:
        return 'Domsheide 16, 28195 Bremen'
    return CASE['loan_id']


def styles():
    return {
        'body': ParagraphStyle('Body', fontName=FONT, fontSize=11, leading=14.5, spaceAfter=9),
        'small': ParagraphStyle('Small', fontName=FONT, fontSize=9, leading=11.5, spaceAfter=6),
        'title': ParagraphStyle('Title', fontName=FONT+'Bold', fontSize=14, leading=17, spaceAfter=12),
        'head': ParagraphStyle('Head', fontName=FONT+'Bold', fontSize=12, leading=15, spaceAfter=10),
    }


def paragraph(text, style):
    return Paragraph(escape(str(text)).replace('\n', '<br/>'), style)


def pdf_document(document):
    target = OUT / document['folder'] / (document['id'] + '.pdf')
    target.parent.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(str(target), pagesize=A4, invariant=1)
    canvas.setAuthor('Klotzkette')
    canvas.setTitle(document['title'])
    canvas.setCreator('Klotzkette')
    st = styles()
    for number, page in enumerate(document['pages'], 1):
        canvas.setFont(FONT+'Bold', 10)
        header = paragraph(document['author'], st['head'])
        _, height = header.wrap(WIDTH, 40)
        header.drawOn(canvas, 52, A4[1] - 35 - height)
        canvas.setFont(FONT, 8.5)
        canvas.drawString(52, A4[1]-72, address(document['author']))
        canvas.setStrokeColor(colors.HexColor('#929292'))
        canvas.line(52, A4[1]-83, A4[0]-52, A4[1]-83)
        y = A4[1]-99
        content = [paragraph(f"{document['date']} | {document['reference']}", st['small'])]
        if number == 1:
            content.extend([paragraph(document['recipient'], st['small']), paragraph(document['title'], st['title'])])
        content.append(paragraph(page['heading'], st['head']))
        content.extend(paragraph(text, st['body']) for text in page.get('paragraphs', []))
        rows = table_data(page)
        if rows:
            widths = page.get('table', {}).get('widths')
            widths = [WIDTH*x/sum(widths) for x in widths] if widths else [WIDTH/len(rows[0])]*len(rows[0])
            t = Table([[paragraph(cell, st['small']) for cell in row] for row in rows], colWidths=widths)
            t.setStyle(TableStyle([
                ('VALIGN', (0,0), (-1,-1), 'TOP'), ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#eceeed')),
                ('LINEBELOW', (0,0), (-1,0), .5, colors.grey), ('LINEBELOW', (0,1), (-1,-1), .2, colors.lightgrey),
                ('LEFTPADDING', (0,0), (-1,-1), 5), ('RIGHTPADDING', (0,0), (-1,-1), 5),
                ('TOPPADDING', (0,0), (-1,-1), 4), ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ]))
            content.append(t)
        content.extend(paragraph(text, st['body']) for text in page.get('after', []))
        for flowable in content:
            _, height = flowable.wrap(WIDTH, A4[1])
            if y - height < 52:
                raise ValueError(f"Seitenüberlauf: {document['id']} Seite {number}, noch {y-52:.1f}, benötigt {height:.1f}")
            flowable.drawOn(canvas, 52, y-height)
            y -= height + (flowable.getSpaceAfter() if hasattr(flowable, 'getSpaceAfter') else 10)
        canvas.setFont(FONT, 8)
        canvas.setFillColor(colors.black)
        canvas.drawString(52, 31, document['reference'])
        canvas.drawRightString(A4[0]-52, 31, f"{number} / {len(document['pages'])}")
        canvas.showPage()
    canvas.save()
    assert len(PdfReader(target).pages) == len(document['pages'])
    return target


def docx_document(document):
    target = OUT / document['folder'] / (document['id'] + '.docx')
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.8), Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(1.8)
    for style in doc.styles:
        if style.type == 1:
            style.font.name = 'Times New Roman'
            style.font.size = Pt(11)
            style.font.color.rgb = RGBColor(0,0,0)
            style.paragraph_format.space_after = Pt(8)
            style.paragraph_format.line_spacing = 1.12
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if key.endswith('Theme'):
                    del fonts.attrib[key]
            for key in ('ascii','hAnsi','eastAsia','cs'):
                fonts.set(qn('w:'+key), 'Times New Roman')
            for border in style.element.xpath('./w:pPr/w:pBdr'):
                border.getparent().remove(border)
    doc.styles['Title'].font.size = Pt(15)
    doc.styles['Heading 1'].font.size = Pt(12)
    doc.styles['Heading 1'].font.bold = True
    header = sec.header.paragraphs[0]
    header.text = document['author']+'\n'+address(document['author'])
    for run in header.runs:
        run.font.size = Pt(9)
    foot = sec.footer.paragraphs[0]
    foot.text = document['reference']+' | Seite '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    foot._p.append(field)
    cp = doc.core_properties
    cp.author = cp.last_modified_by = 'Klotzkette'
    cp.title = document['title']
    cp.language = 'en-GB' if 'Agreement' in document['title'] else 'de-DE'
    cp.created = cp.modified = datetime.fromisoformat(document['date'])
    for i, page in enumerate(document['pages']):
        if i:
            doc.add_page_break()
        else:
            doc.add_paragraph(document['recipient'])
            doc.add_paragraph(document['date']+' | '+document['reference'])
            doc.add_paragraph(document['title'], 'Title')
        doc.add_paragraph(page['heading'], 'Heading 1')
        for text in page.get('paragraphs', []):
            doc.add_paragraph(text)
        rows = table_data(page)
        if rows:
            t = doc.add_table(rows=0, cols=len(rows[0]))
            t.style = 'Table Grid'
            for row in rows:
                cells = t.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)
        for text in page.get('after', []):
            doc.add_paragraph(text)
    doc.save(target)
    return target


def bank_records():
    balances = 84300
    pages, ledger = [], []
    # Laufende Kontobewegungen, keine betriebswirtschaftliche Auswertung.
    customer_receipts = [164900,158700,149100,177200,169400,184500,173800,159600,181900,167700,173300,169500,178200,165400,171800,167500,155800,598600]
    for index, payment_date in enumerate(PAYMENTS):
        year, month, _ = map(int, payment_date.split('-'))
        ym = f'{year}-{month:02d}'
        previous = balances
        income = customer_receipts[index]
        entries = [
            (f'{ym}-02', 'Nordhafen Antriebstechnik GmbH', 'Sammelzahlung Ventilblöcke, Projekt N42', income),
            (f'{ym}-04', 'Hanseblech Zuschnitt GmbH', f'Werkstofflieferung HB-{year}-{month:02d}', -21300),
            (f'{ym}-06', 'Werftkomponenten Kroll KG', f'Messgehäuse WK-{year}-{month:02d}', 26800),
            (f'{ym}-08', 'Gewerbehof Hüttenstraße KG', 'Miete Halle 4 und Nebenkosten', -9200),
            (payment_date, 'Vellio Finance GmbH', f'NF-WP-221014-01 Zinsperiode {index+1:02d}', -5000),
        ]
        if index == 17:
            entries.append((payment_date, 'Vellio Finance GmbH', 'NF-WP-221014-01 Endtilgung', -500000))
        entries.extend([
            (f'{ym}-18', 'Prüftechnik Bruns GmbH', 'Dichtheitsprüfungen und Materialzeugnisse', -8400),
            (f'{ym}-20', 'Versorgungswerke Weserbogen GmbH', 'Strom- und Wärmeabschlag Werkhalle', -4700),
            (f'{ym}-22', 'Finanzkasse Bremen', 'Abgaben, Buchungsreferenz 72-481', -18600),
            (f'{ym}-25', 'Beschäftigte Weserfunken', 'Sammelüberweisung Löhne und Gehälter', -64200),
            (f'{ym}-27', 'Sozialkassen', 'Beitragsnachweise laufender Monat', -27900),
            (f'{ym}-28', 'Werkzeugservice Evers GmbH', 'Fräser, Bohrer und Instandsetzung Spindel', -16200),
        ])
        if index >= 14:
            entries.append((f'{ym}-28', 'Altmetall Nordufer GmbH', 'Ankauf Späne und Reststücke', 2100))
        rows = []
        for booking_date, recipient, purpose, amount in entries:
            balances += amount
            rows.append([booking_date[8:]+'.'+booking_date[5:7]+'.', recipient+'\n'+purpose, money(amount), money(balances)])
            ledger.append([booking_date, recipient, purpose, str(amount), str(balances), f'K4 Blatt {index+1}'])
        pages.append({
            'heading': f'{index+1}. Kontoauszug {month:02d}/{year}',
            'paragraphs': [f'Kontoinhaber: Weserfunken Präzisionsbau GmbH, Hüttenstraße 43, 28237 Bremen. Geschäftskonto 01728400, Währung EUR. Vortrag aus dem vorherigen Auszug: {money(previous)} EUR. Elektronisch abgerufen von Jutta Kröger am 06.06.2026, 09:{index+10:02d} Uhr.', 'Buchungstag und Wertstellung stimmen bei den aufgeführten Überweisungen überein. Empfängerangaben stammen aus dem gespeicherten Zahlungsauftrag. Zinszahlungen an Vellio Finance GmbH gehen an AT61 1904 3002 3457 3201.'],
            'table': {'headers':['Datum','Buchung / Verwendungszweck','Betrag EUR','Saldo EUR'], 'rows': rows, 'widths':[.7,3.7,1.15,1.15]},
            'after': [f'Neuer Saldo: {money(balances)} EUR. Die Darlehenszahlung ist im SEPA-Einzelauftrag gespeichert; Sammelüberweisungen betreffen ausschließlich die gesondert bezeichneten Zahlungen. Ansprechpartner für Auszugsduplikate: Sven Peters, Geschäftskundenservice, service@weserkontor.example.'],
        })
    return {'id':'04_K4_Kontoauszuege','folder':'01_eingang','title':'Anlage K4 | Geschäftskonto, November 2022 bis April 2024','author':'Weserkontor Bank eG','recipient':CASE['debtor']['name'],'date':'2026-06-06','reference':'Konto 01728400 / NF-WP-221014-01','pages':pages}, ledger


def other_records():
    return [
        {'id':'03_K3_Auszahlung','folder':'01_eingang','title':'Anlage K3 | Ausführungsbestätigung','author':CASE['bank']['name'],'recipient':CASE['debtor']['name'],'date':'2022-10-14','reference':CASE['loan_id'],'pages':[{'heading':'1. EUR-Zahlung ausgeführt','paragraphs':[
            'Sehr geehrter Herr Grothe, die Auszahlung des Darlehens ist heute um 10:11 Uhr ausgelöst und mit Wertstellung 14.10.2022 bestätigt worden. Der Auszahlungsauftrag wurde nicht um ein Disagio oder eine Bearbeitungsgebühr gekürzt. Der überwiesene Betrag beträgt 500.000,00 EUR.',
            'Auftraggeber ist Nexora Frontbank AG. Begünstigter ist Weserfunken Präzisionsbau GmbH. Die empfangende Bank führt das Geschäftskonto 01728400. Als End-to-End-Referenz wurde NF-WP-221014-01-DRAW verwendet. Die Zahlungsbestätigung enthält ausschließlich diesen Auszahlungsvorgang; sie ist keine Saldenbestätigung für die weitere Vertragslaufzeit.',
            'Die noch heute vorgesehene Übernahme wird separat bestätigt. Bis zum Wirksamkeitszeitpunkt erreichen Sie unsere Kreditabwicklung unter legal@nexora.example. Für die Unterzeichnung durch alle drei Beteiligten wird Ihnen die unterschriebene Fassung des Transfer and Assumption Agreement nach Abschluss der letzten Signatur zugeleitet.',
            'Bitte reichen Sie diese Bestätigung an Frau Kröger weiter. Die Buchhaltung benötigt für die Zuordnung nicht nur den Kontoauszug, sondern auch die Darlehensnummer. Ein abweichender Verwendungszweck bei späteren Rückzahlungen verzögert die Zuordnung.',
            'Mit freundlichen Grüßen\nDr. Falk Neubauer\nNexora Frontbank AG\nKreditabwicklung Wien',
        ]}]},
        {'id':'00_Gerichtliche_Verfuegung','folder':'01_eingang','title':'Schriftliches Vorverfahren und Zustellungsunterlagen','author':'Landgericht Bremen, 4. Zivilkammer','recipient':'Nexora Frontbank AG, Schottenring 18, 1010 Wien, Österreich','date':'2026-09-10','reference':CASE['docket'],'pages':[{'heading':'1. Verfügung','paragraphs':[
            'In dem Rechtsstreit Kunigunde von Tettenborn als Insolvenzverwalterin über das Vermögen der Weserfunken Präzisionsbau GmbH gegen Nexora Frontbank AG wird das schriftliche Vorverfahren angeordnet. Die Klageschrift vom 08.09.2026 nebst den Anlagen K1 bis K12 ist der Beklagten zuzustellen.',
            'Die Beklagte wird aufgefordert, binnen einer Notfrist von zwei Wochen nach Zustellung der Klageschrift durch einen Rechtsanwalt anzuzeigen, ob sie sich gegen die Klage verteidigen will. Die Frist kann nicht verlängert werden. Vor dem Landgericht besteht Anwaltszwang. Eine persönliche Mitteilung der Partei ersetzt die anwaltliche Verteidigungsanzeige nicht.',
            'Für die schriftliche Klageerwiderung wird eine weitere Frist von drei Wochen nach Ablauf der Notfrist gesetzt. Die Beklagte soll ihre Einwendungen, Beweismittel und etwaige Zuständigkeitsrügen in der Erwiderung darlegen. Der Eingang bei Gericht ist maßgeblich. Verspäteter Vortrag kann nach den gesetzlichen Voraussetzungen zurückgewiesen werden.',
            'Die Beklagte wird darauf hingewiesen, dass bei unterbliebener rechtzeitiger Verteidigungsanzeige auf Antrag ohne mündliche Verhandlung ein Versäumnisurteil ergehen kann. Die Klägerin hat einen entsprechenden Antrag gestellt. Die weitergehenden Voraussetzungen werden vom Gericht geprüft.',
            'Die elektronische Einreichung anwaltlicher Schriftsätze richtet sich nach den Vorschriften der Zivilprozessordnung. Anlagen sind eindeutig zu bezeichnen. Das gerichtliche Aktenzeichen ist bei jeder Eingabe anzugeben. Ein Termin zur mündlichen Verhandlung wird derzeit nicht bestimmt.',
            'gez. Dr. Merz, Richter am Landgericht\nFür die Geschäftsstelle: S. Hagedorn, Justizfachangestellter\nAbschrift zur Zustellung, ausgefertigt am 10.09.2026',
        ]},{'heading':'2. Rücklauf Zustellung','paragraphs':[
            'Zustellungsbezug: Landgericht Bremen, 4 O 1186/26. Empfänger: Nexora Frontbank AG, Schottenring 18, 1010 Wien, Österreich. Übermittelt wurden die Verfügung vom 10.09.2026, die Klageschrift vom 08.09.2026 und die Anlagen K1 bis K12.',
            'Übergabe an der Geschäftsanschrift am Dienstag, 15.09.2026, um 10:24 Uhr. Entgegengenommen durch Lena Berger, Empfang. Die Sendung wurde noch am selben Vormittag an Dr. Falk Neubauer weitergegeben. Eingangsstempel der Rechtsabteilung: 15.09.2026, 11:02 Uhr.',
            'Der Empfangsbeleg und die Umschlagvorderseite wurden am 16.09.2026 um 08:46 Uhr als gemeinsamer Scan an die Kanzlei Rademacher Westhoff übermittelt. Die bei der Bank aufbewahrte Sendung enthält 100 Blatt Klage und Anlagen sowie die gerichtliche Verfügung. Auf dem Umschlag befindet sich keine andere Zustellanschrift.',
            'Bearbeitungsvermerk Empfang: Der zunächst angekündigte Kurier für den Quartalsabschluss war eine andere Sendung. Die Gerichtsunterlagen wurden nicht mit dieser Post zusammengeführt. Rückfragen zur Entgegennahme an empfang@nexora.example, Lena Berger.',
            'Abschrift des Eingangsvermerks: L. Berger, 16.09.2026. Die Unterschrift auf dem Papierbeleg verbleibt in der Wiener Poststelle.',
        ]}]},
        {'id':'00_Verteidigungsanzeige_20260922','folder':'03_korrespondenz','title':'Verteidigungsanzeige','author':CASE['defence_firm']['name'],'recipient':'Landgericht Bremen, Domsheide 16, 28195 Bremen','date':'2026-09-22','reference':CASE['docket'],'pages':[{'heading':'1. Anzeige der Verteidigungsbereitschaft','paragraphs':[
            'In dem Rechtsstreit Kunigunde von Tettenborn als Insolvenzverwalterin gegen Nexora Frontbank AG zeigen wir die Vertretung der Beklagten an. Die Beklagte wird sich gegen die Klage verteidigen. Die Klageschrift ist der Beklagten am 15.09.2026 zugestellt worden.',
            'Die Vertretung erstreckt sich auf die Nexora Frontbank AG, nicht auf die Vellio Finance GmbH. Eine Vollmacht zur Vertretung der letztgenannten Gesellschaft liegt uns nicht vor. Wir bitten, diese Unterscheidung bei der Bezeichnung der Parteien und der Zustellung künftiger Verfügungen zu berücksichtigen.',
            'Eine Einlassung zur Hauptsache ist mit dieser Anzeige nicht verbunden. Insbesondere bleibt die Rüge der internationalen Zuständigkeit hinsichtlich der neben der Insolvenzanfechtung geltend gemachten Ansprüche vorbehalten. Eine Zustimmung zu einer Zuständigkeit kraft Einlassung wird nicht erklärt.',
            'Die Klageerwiderung erfolgt innerhalb der gesetzten Frist. Wir bitten um elektronische Zustellung an das Anwaltspostfach des Unterzeichners. Unser Aktenzeichen lautet 26-184-NF.',
            'Dr. Henrik Westhoff\nRechtsanwalt\nElektronisch übermittelt am 22.09.2026, 14:18 Uhr; Eingangsbestätigung 14:18:07 Uhr in der Handakte.',
        ]}]},
    ]


def main():
    for suffix, bold in [('', False), ('Bold', True)]:
        pdfmetrics.registerFont(TTFont(FONT+suffix, str(serif_font_path(bold))))
    documents, emails, notes = [], [], []
    for name in ['claim', 'contracts', 'defence', 'correspondence', 'business-records']:
        data = json.loads((FIXTURES / (name+'.json')).read_text())
        documents.extend(data['documents'])
        emails.extend(data.get('emails', []))
        notes.extend(data.get('notes', []))
    bank, ledger = bank_records()
    documents.extend([bank, *other_records()])
    exhibits = [d for d in documents if re.match(r'\d+_K\d+_', d['id'])]
    if len(documents) != 22 or len(exhibits) != 12 or sum(len(d['pages']) for d in exhibits) != 75:
        raise ValueError('Unvollständiger Quellenbestand: 22 Dokumente und 12 K-Anlagen mit 75 Seiten erforderlich')
    if len(emails) != 10 or len(notes) != 2:
        raise ValueError('Korrespondenz nicht vollständig')
    if any('§' in json.dumps(d, ensure_ascii=False) for d in documents):
        raise ValueError('Paragrafzeichen in Quelldokumenten; vor der Ausgabe redaktionell korrigieren')
    seen = set()
    manifest = []
    for document in documents:
        key = document['folder']+'/'+document['id']
        if key.casefold() in seen:
            raise ValueError('Doppeltes Dokument: '+key)
        seen.add(key.casefold())
        native = document['id'].startswith(('00_Klageerwiderung', '01_B1_'))
        path = docx_document(document) if native else pdf_document(document)
        manifest.append({'path':str(path.relative_to(OUT)), 'pages':len(document['pages']), 'title':document['title']})
    for mail in emails:
        msg = EmailMessage(policy=policy.SMTP)
        for key in ('from','to','cc'):
            if mail.get(key):
                msg[key.title()] = mail[key]
        msg['Date'] = datetime.fromisoformat(mail['date'])
        msg['Subject'] = mail['subject']
        msg['Message-ID'] = f"<{mail['id']}@weserfunken.example>"
        msg.set_content(mail['body'])
        target = OUT/mail['folder']/(mail['id']+'.eml')
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(bytes(msg))
    for note in notes:
        target = OUT/note['folder']/(note['id']+'.txt')
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(note['title']+'\n\n'+note['text']+'\n', encoding='utf-8')
    csv_path = OUT/'03_korrespondenz/30_Buchungsdaten_Geschaeftskonto.csv'
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.writer(handle, delimiter=';')
        writer.writerow(['Buchungstag','Empfänger','Verwendungszweck','Betrag EUR','Saldo EUR','Beleg'])
        writer.writerows(ledger)
    manifest_path = FIXTURES/'documents.json'
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+'\n')
    readme = OUT/'README.md'
    if readme.exists():
        begin, end = '<!-- BEGIN fintech-inventory -->', '<!-- END fintech-inventory -->'
        lines = [begin, '## 1.5. Einzelunterlagen', '', NOTICE_MARKDOWN, '', '| Datei | Inhalt |', '| --- | --- |']
        for row in sorted(manifest, key=lambda item: item['path']):
            label = row['title'].replace('|', ' / ')
            lines.append(f"| [{row['path']}]({row['path']}) | {label} |")
        for mail in emails:
            path = mail['folder']+'/'+mail['id']+'.eml'
            lines.append(f"| [{path}]({path}) | {mail['subject']} |")
        for note in notes:
            path = note['folder']+'/'+note['id']+'.txt'
            lines.append(f"| [{path}]({path}) | {note['title']} |")
        lines.append('| [30_Buchungsdaten_Geschaeftskonto.csv](03_korrespondenz/30_Buchungsdaten_Geschaeftskonto.csv) | Kontobuchungen mit Salden und Belegbezug |')
        lines.extend(['', end, ''])
        block = '\n'.join(lines)
        text = readme.read_text()
        text = re.sub(re.escape(begin)+'.*?'+re.escape(end), lambda _: block.rstrip(), text, flags=re.S) if begin in text else text.rstrip()+'\n\n'+block
        readme.write_text(text, encoding='utf-8')
    print(f'{len(documents)} Dokumente, {len(emails)} E-Mails, {len(notes)} Notizen, {len(ledger)} Kontobuchungen')


if __name__ == '__main__':
    main()
