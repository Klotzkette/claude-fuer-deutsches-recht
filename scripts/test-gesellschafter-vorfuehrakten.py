#!/usr/bin/env python3
"""Regressionsprüfungen der zwei unabhängigen Gesellschaftsrechtsakten."""
import csv
from datetime import date
from email import policy
from email.parser import BytesParser
import hashlib
import io
import os
from pathlib import Path
import re
import unittest
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from vorfuehrakte_gesellschafter_berlin import CASE as BERLIN
from vorfuehrakte_gesellschafter_muenchen import CASE as MUENCHEN
from testakte_disclaimer import NOTICE_DE, NOTICE_EN, NOTICE_FILENAME

ROOT = Path(__file__).resolve().parents[1]
CASES = [BERLIN, MUENCHEN]
TABLES = {BERLIN['slug']:'19_Projektkonto.xlsx', MUENCHEN['slug']:'18_Beteiligungsrechnung.xlsx'}


def directory(case):
    return ROOT / 'testakten' / case['slug']


class CasesTest(unittest.TestCase):
    def test_exact_source_inventory(self):
        for case in CASES:
            expected = {d['file'] for d in case['documents']} | {TABLES[case['slug']]}
            actual = {p.name for p in directory(case).iterdir() if p.is_file() and p.name not in {'README.md','rubric.yaml'}}
            self.assertEqual(expected, actual, case['slug'])

    def test_distinct_source_formats(self):
        for case in CASES:
            extensions = {p.suffix for p in directory(case).glob('[0-9]*')}
            self.assertTrue({'.docx','.eml','.pdf','.csv','.txt','.xlsx'} <= extensions)

    def test_emails_have_complete_headers_and_exact_attachments(self):
        for case in CASES:
            for item in case['documents']:
                if not item['file'].endswith('.eml'):
                    continue
                message = BytesParser(policy=policy.default).parsebytes((directory(case)/item['file']).read_bytes())
                for header in ['From','To','Subject','Date','Message-ID']:
                    self.assertTrue(message[header], (item['file'],header))
                    self.assertFalse(message[header].defects,(item['file'],header,message[header].defects))
                self.assertEqual(message.get_body(preferencelist=('plain',)).get_content().replace('\r\n','\n').strip(), item['body'])
                attachments = {part.get_filename():part.get_payload(decode=True) for part in message.iter_attachments()}
                self.assertEqual(set(attachments), set(case['attachments'].get(item['file'],[])))
                for name, content in attachments.items():
                    self.assertEqual(content, (directory(case)/name).read_bytes(), name)

    def test_investor_message_reaches_both_founders(self):
        path=directory(MUENCHEN)/'04_Angebot_Investor.eml'
        message=BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        self.assertEqual({a.addr_spec for a in message['To'].addresses},
                         {'hs@isarwinkel-geraetebau.example','qr@isarwinkel-geraetebau.example'})

    def test_csv_rectangles_and_minimum_rows(self):
        for case in CASES:
            for item in case['documents']:
                if not item['file'].endswith('.csv'):
                    continue
                with (directory(case)/item['file']).open(encoding='utf-8-sig',newline='') as stream:
                    rows=list(csv.reader(stream,delimiter=';'))
                self.assertEqual(rows[0],item['headers'])
                self.assertGreaterEqual(len(rows),5)
                self.assertTrue(all(len(r)==len(rows[0]) for r in rows))

    def test_no_solution_or_training_labels_in_originals(self):
        pattern=re.compile(r'Testakte|fiktiv|Platzhalter|Formathinweis|Musterlösung|Lösungsmatrix|Red-Team|§',re.I)
        for case in CASES:
            for item in case['documents']:
                self.assertIsNone(pattern.search(item['title']+'\n'+item['body']),item['file'])

    def test_docx_contains_entire_authored_body(self):
        for case in CASES:
            for item in case['documents']:
                if not item['file'].endswith('.docx'):
                    continue
                text='\n'.join(p.text for p in Document(directory(case)/item['file']).paragraphs)
                for paragraph in re.split(r'\n\s*\n',item['body']):
                    self.assertIn(paragraph.removeprefix('## '),text,item['file'])

    def test_native_pdf_is_readable_and_has_no_warning(self):
        for case in CASES:
            for path in directory(case).glob('*.pdf'):
                reader=PdfReader(path)
                text=' '.join(p.extract_text() for p in reader.pages)
                self.assertGreater(len(text),600,path.name)
                self.assertNotIn(NOTICE_DE,text)
                self.assertNotIn(NOTICE_EN,text)

    def test_berlin_litigation_and_contract_are_distinct(self):
        items={d['file']:d for d in BERLIN['documents']}
        self.assertIn('nicht gekündigt',items['01_Mandat_Rabenstein.eml']['body'])
        self.assertIn('Über den Anstellungsvertrag wird nicht abgestimmt',items['08_Niederschrift_K5.docx']['body'])
        self.assertIn('8750 Stimmen Nein',items['08_Niederschrift_K5.docx']['body'])
        self.assertIn('10000 Stimmen Ja',items['08_Niederschrift_K5.docx']['body'])
        self.assertIn('Paragraf 38 Absatz 2 GmbHG',items['02_Klage_Seidel.docx']['body'])

    def test_service_dates_leave_work_time_at_case_date(self):
        self.assertEqual((date(2026,10,9)-date(2026,9,25)).days,14)
        self.assertEqual((date(2026,10,23)-date(2026,10,9)).days,14)
        self.assertLess(date(2026,10,2),date(2026,10,9))
        text=(directory(BERLIN)/'04_Zustellung_Posteingang.txt').read_text()
        self.assertIn('25.09.2026',text)
        self.assertIn('32 O 187/26',text)

    def test_court_notice_explains_costs_and_enforcement(self):
        item=next(d for d in BERLIN['documents'] if d['file']=='03_Gerichtliche_Verfuegung.pdf')
        for phrase in ['notwendigen Kosten der Gegenseite','Paragraf 91 ZPO',
                       'Paragraf 708 Nummer 2 ZPO','ohne Sicherheitsleistung',
                       'Notfrist ist nicht verlängerbar']:
            self.assertIn(phrase,item['body'])

    def test_berlin_bank_matches_csv(self):
        case=directory(BERLIN)
        wb=load_workbook(case/TABLES[BERLIN['slug']],data_only=True)
        s=wb['Projektkonto']
        self.assertEqual(s['F14'].value,28200)
        self.assertEqual(s['D14'].value,37000)
        self.assertEqual(s['E14'].value,40800)
        with (case/'17_Bankprotokoll.csv').open(encoding='utf-8-sig',newline='') as stream:
            rows=list(csv.DictReader(stream,delimiter=';'))
        for r,row in enumerate(rows,7):
            self.assertEqual(s.cell(r,2).value,row['Referenz'])
            self.assertEqual(s.cell(r,4).value,float(row['Eingang_EUR']))
            self.assertEqual(s.cell(r,5).value,float(row['Ausgang_EUR']))
        self.assertIn('yy',s['A7'].number_format)

    def test_munich_capital_and_premium(self):
        wb=load_workbook(directory(MUENCHEN)/TABLES[MUENCHEN['slug']],data_only=True)
        s=wb['Kapitalaufnahme']
        for cell,value in {'E9':31250,'F6':0.48,'F7':0.32,'F8':0.2,'B12':200000,'B13':6250,'B14':193750,'B15':0}.items():
            self.assertAlmostEqual(s[cell].value,value,places=8)

    def test_munich_later_payment_is_a_later_filing_note(self):
        item=next(d for d in MUENCHEN['documents'] if d['file']=='03_Gesellschafterliste.docx')
        original, note=item['body'].split('Ablagevermerk 30.09.2026',1)
        self.assertNotIn('19.07.2021',original)
        self.assertIn('19.07.2021',note)

    def test_workbooks_have_live_formulas_and_no_errors(self):
        for case in CASES:
            path=directory(case)/TABLES[case['slug']]
            wb=load_workbook(path,data_only=False)
            self.assertGreater(sum(c.data_type=='f' for s in wb for row in s for c in row),5)
            cached=load_workbook(path,data_only=True)
            for s in cached:
                for row in s:
                    for cell in row:
                        self.assertNotEqual(cell.data_type,'e',(case['slug'],cell.coordinate,cell.value))

    def test_readme_core_set_exists_and_is_small(self):
        for case in CASES:
            text=(directory(case)/'README.md').read_text()
            self.assertLessEqual(len(case['core']),8)
            self.assertIn('60-Minuten',text)
            for name in case['core']:
                self.assertTrue((directory(case)/name).is_file())
                self.assertIn(f']({name})',text)
            self.assertIn('reserved-example-contacts',text)
            self.assertIn(NOTICE_DE,text)
            self.assertIn(NOTICE_EN,text)

    def test_combined_pdf_has_all_document_bookmarks(self):
        for case in CASES:
            reader=PdfReader(directory(case)/'gesamt-pdf'/f"{case['slug']}_gesamt.pdf")
            text='\n'.join(p.extract_text() for p in reader.pages)
            self.assertNotIn(NOTICE_DE,text)
            self.assertNotIn(NOTICE_EN,text)
            self.assertGreaterEqual(len(reader.pages),len(case['documents'])+1)
            outline=str(reader.outline)
            for item in case['documents']:
                self.assertIn(Path(item['file']).stem,outline,item['file'])
            self.assertIn(Path(TABLES[case['slug']]).stem,outline)

    @unittest.skipUnless(os.environ.get('VORFUEHRAKTEN_ZIPS'),'VORFUEHRAKTEN_ZIPS für Archivprüfung setzen')
    def test_flat_archives_with_real_files_and_disclaimer(self):
        for case in CASES:
            native={d['file'] for d in case['documents']}|{TABLES[case['slug']]}
            for suffix in ('','-einzelpdfs'):
                path=Path(os.environ['VORFUEHRAKTEN_ZIPS'])/f"testakte-{case['slug']}{suffix}.zip"
                with zipfile.ZipFile(path) as z:
                    self.assertIsNone(z.testzip())
                    names=z.namelist()
                    self.assertEqual(len(names),len(set(names)))
                    self.assertTrue(all('/' not in n and not n.endswith('.md') for n in names))
                    notice=z.read(NOTICE_FILENAME).decode('utf-8')
                    self.assertIn(NOTICE_DE,notice);self.assertIn(NOTICE_EN,notice)
                    if suffix:
                        self.assertEqual(len(names),len(native)+1)
                        for name in names:
                            if name!=NOTICE_FILENAME:
                                self.assertTrue(name.endswith('.pdf'))
                                pages=PdfReader(io.BytesIO(z.read(name))).pages
                                self.assertGreater(len(pages),0)
                                content='\n'.join(p.extract_text() for p in pages)
                                self.assertNotIn(NOTICE_DE,content)
                                self.assertNotIn(NOTICE_EN,content)
                    else:
                        combined=f"{case['slug']}_gesamt.pdf"
                        self.assertEqual(native|{NOTICE_FILENAME,combined},set(names))
                        self.assertEqual(z.read(combined),(directory(case)/'gesamt-pdf'/combined).read_bytes())
                        for name in native:
                            self.assertEqual(hashlib.sha256(z.read(name)).digest(),hashlib.sha256((directory(case)/name).read_bytes()).digest())
                    self.assertFalse(any('rubric' in n or 'README.md'==n for n in names))


if __name__=='__main__':
    unittest.main()
