#!/usr/bin/env python3
"""Fallbezogene Ergänzungen 11 bis 26; Quellen 01 bis 10 bleiben unverändert.

Aufruf: python scripts/build-schulstrasse-nutzungsbelege.py [--check-only]
Optional: --upload-zip PFAD prüft die zehn Originale zusätzlich gegen den Upload.
Gesamt-PDF, Archive und zentrale Metadaten sind nicht Aufgabe dieses Builders.

Bestandsprüfung am 22.09.2026: Alle zehn Originale wurden byteweise verglichen
und vollständig gelesen: DOCX einschließlich Kopf-/Fußzeilen, beide XLSX-Blätter,
EML-Header und decodierte Texte, CSV sowie beide 17-seitigen Gesamt-PDFs.
Upload-Gesamt-PDF: SHA-256
3f194e925136be334bcc709c0418362f33da82bc058cbc7ca41ee6136f0677de.
Sie wich von der damaligen Repo-Lesefassung ab, insbesondere bei der Extraktion
der Office-Druckseiten. Die zehn nativen Originale waren sämtlich bytegleich.
Das Upload-README.txt ist Archiv-Metadatum, kein weiteres Aktenstück.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
from datetime import datetime, timezone
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "testakten/strassenverkehrsrecht-stvo-schulstrasse-lieferzone"
ORIGINALS = {
    "01_verkehrsanordnung_schulstrasse_buchenweg.docx": "ed855b9d5285beb8bf5e70300e99a1481eda457d79c716abe22ab0a033a7351b",
    "02_lieferfahrten_und_ausnahmegenehmigungen.xlsx": "7b63af6312c4b9e633eeb4a63912d6d1867d06883367b225bb9cde24e1397832",
    "03_email_baeckerei_ausnahmegenehmigung_2026-08-14.eml": "9cf932bccfdd0d10fcb7ea665389e4e17c6d7328c3969a2d7aedaa844ea62cb7",
    "04_bussgeldanhoerung_lieferwagen_2026-08-20.docx": "3240f4c945014aa3dcff1ee7e938e4d4c32efe5b227ec563e19cd83203bdb109",
    "05_verkehrszaehlung_buchenweg_roh.csv": "9cb9a96e28861465138f504363d88c00ce8d7856849ec1f7d23c3d48712197de",
    "06_email_elternvertretung_beobachtungen_2026-08-18.eml": "10b082ea55a666f491e12c7fb39d0f1ce34de3ab0600c2d4ec82f5c9bf77e447",
    "07_fahrernotiz_lieferung_2026-08-19.docx": "f1d26ea0bff655555da14aeaf9b219b57cda61a9e63690a8d15edb526203cbf7",
    "08_telefonvermerk_maessner_kruse_2026-08-21.docx": "d898e8d6ec675243c79615a13b8bd0bffcffac592e75b02cfb0588fdf60b025a",
    "09_email_disposition_intern_2026-08-19.eml": "ce2324f9b0cfd1db6d51a4701645fa6cecbf61a01e9e19e4b4c1970cec5f92ad",
    "10_zwischennachricht_stadt_ausnahmeantrag_2026-08-25.docx": "dee106d482daa58e6290d74079393fcc17a83b47d881b5f8e130582f99ae52f4",
}
FILES = {
    11: "11_ergaenzung_zufahrtsantrag_2026-08-26.docx",
    12: "12_nachreichung_fahrzeugdaten_2026-08-24.txt",
    13: "13_lieferschein_2026-08-19-031.pdf",
    14: "14_verkehrszeichenplan_va884_blatt2_abschrift.pdf",
    15: "15_strassenbestandsakte_widmung_2026-08-26.docx",
    16: "16_email_disposition_planberichtigung_2026-08-26.eml",
    17: "17_kontrollprotokoll_auszug_2026-08-26.docx",
    18: "18_email_stadt_ermittlungsfragen_2026-08-26.eml",
    19: "19_ortstermin_lindenplatz_2026-08-27.docx",
    20: "20_massskizze_ladeweg_poller_2026-08-27.pdf",
    21: "21_ladezeiten_filiale_2026-08-25_bis_28.csv",
    22: "22_email_tiefbau_container_poller_2026-08-28.eml",
    23: "23_filialnotiz_ladeversuche_2026-08-28.txt",
    24: "24_antrag_rollbox_stellflaeche_2026-08-31.docx",
    25: "25_email_schule_lieferfenster_2026-08-31.eml",
    26: "26_stadt_sachstand_nutzungsantraege_2026-09-01.docx",
}
FIXED = datetime(2026, 9, 1, 12, tzinfo=timezone.utc)
CITY = ("STADT EICHENSTEDT", "Straßenverkehrsbehörde", "Marktplatz 1, 38118 Eichenstedt", "Telefon 0531 771 362 | verkehr@eichenstedt.de")
BAKERY = ("BÄCKEREI KRUSE & SOHN KG", "Geschäftsführung und Disposition", "Mühlenstraße 7, 38118 Eichenstedt", "Telefon 0531 882 117 | h.kruse@kruse-brot.de")


def verify_originals(upload: Path | None = None) -> None:
    for name, digest in ORIGINALS.items():
        if hashlib.sha256((CASE / name).read_bytes()).hexdigest() != digest:
            raise ValueError(f"Original verändert; kein automatisches Überschreiben: {name}")
    if upload:
        with ZipFile(upload) as archive:
            for name in ORIGINALS:
                if archive.read(name) != (CASE / name).read_bytes():
                    raise ValueError(f"Upload weicht ab: {name}")


def save_doc(doc: Document, number: int) -> None:
    # Word-Vorlagen enthalten teils Themenfonts und eine geerbte Titellinie.
    for root in (doc.styles.element, doc.element):
        for border in root.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        for fonts in root.xpath(".//w:rFonts"):
            for key in list(fonts.attrib):
                if "theme" in key.lower():
                    del fonts.attrib[key]
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + key), "Times New Roman")
    # Feste OOXML-ZIP-Zeitstempel machen Wiederholungsbauten bytegleich.
    raw = io.BytesIO()
    doc.save(raw)
    with ZipFile(raw) as src, ZipFile(CASE / FILES[number], "w", ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info = ZipInfo(name, (2026, 9, 1, 12, 0, 0))
            info.compress_type = ZIP_DEFLATED
            dst.writestr(info, src.read(name))


def letter(number: int, sender: tuple[str, ...], recipient: str, date: str,
           reference: str, title: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.1)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(14)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(9)
    doc.core_properties.author = sender[0]
    doc.core_properties.title = title
    doc.core_properties.created = FIXED
    doc.core_properties.modified = FIXED
    p = doc.add_paragraph()
    r = p.add_run(sender[0]); r.bold = True; r.font.size = Pt(14)
    p.paragraph_format.space_after = Pt(3)
    p = doc.add_paragraph("\n".join(sender[1:]))
    p.paragraph_format.space_after = Pt(18)
    doc.add_paragraph(recipient)
    doc.add_paragraph(f"Eichenstedt, {date}\nGeschäftszeichen: {reference}")
    doc.add_paragraph(title, "Title")
    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.space_before = Pt(0)
    footer.add_run(f"{reference} | Seite ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def body(doc: Document, *paragraphs: str) -> None:
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)


def heading(doc: Document, title: str) -> None:
    doc.add_paragraph(title, "Heading 1")


def email(number: int, sender: str, recipient: str, cc: str, date: str,
          subject: str, content: str, reply: str | None = None) -> None:
    msg = EmailMessage(policy=policy.SMTP)
    domain = sender.split("@", 1)[1].rstrip(">")
    msg["From"], msg["To"], msg["Cc"] = sender, recipient, cc
    msg["Date"], msg["Subject"] = date, subject
    msg["Message-ID"] = f"<buchenweg-{number}-202608@{domain}>"
    msg["Reply-To"] = sender
    if reply:
        msg["In-Reply-To"], msg["References"] = reply, reply
    msg.set_content(content.strip() + "\n", charset="utf-8", cte="quoted-printable")
    (CASE / FILES[number]).write_bytes(msg.as_bytes())


def create_letters() -> None:
    d = letter(11, BAKERY, "Stadt Eichenstedt\nStraßenverkehrsbehörde\nHerrn Timo Mäßner\nMarktplatz 1\n38118 Eichenstedt", "26. August 2026", "66.2-VA-2026-884-A03", "Ergänzung unseres Zufahrtsantrags")
    body(d, "Sehr geehrter Herr Mäßner,",
         "wir halten den Antrag vom 14. August für die Filiale Buchenweg 18 aufrecht. Beantragt ist an Schultagen montags bis freitags jeweils eine Zufahrt zwischen 07:35 und 07:55 Uhr, ab Erteilung bis zum Ende des Verkehrsversuchs am 18. Dezember 2026. Wir benötigen keine Zufahrt während der mittäglichen Sperrzeit. Die Fahrzeuge BS-KR 418 und BS-KR 422 sollen einander ersetzen, nicht gleichzeitig fahren.",
         "Die Tour fährt vom Lindenplatz zur Filiale und weiter über die Finkenstraße aus. Gehalten werden soll auf der Fahrbahn vor Hausnummer 18, nicht auf dem Gehweg. Die Übergabe dauert bei freiem Zugang gewöhnlich sechs bis neun Minuten. Währenddessen bleibt der Fahrer am Fahrzeug; Nele Maaß übernimmt die Kisten an der Ladentür. Wir können auf eine Rücknahme leerer Kisten am Morgen verzichten und diese am Nachmittag abholen.",
         "Die erste Ladung umfasst gewöhnlich 14 bis 18 Brotkisten und zwei geschlossene Kühlboxen. Für einen Versuch können haltbare Ware und Leergut getrennt gefahren werden. Die frischen Brote aus Nordring sind aber erst nach 07:10 Uhr transportfertig. Die Filiale öffnet um 07:00 Uhr mit dem Restbestand; mehrere Kunden kommen vor dem Schulbeginn und warten derzeit auf die zweite Auslage.",
         "Bitte beziehen Sie in den Ortstermin auch die drei Poller am Gehwegzugang vom Lindenplatz ein. Ein mittlerer Poller wirkt herausnehmbar. Unsere Sackkarre kommt einzeln durch, Begegnungen mit Kinderwagen sind dort schwierig. Wir bitten um Auskunft, ob zeitweise eine breitere Passage hergestellt werden kann und wer diese bedienen dürfte. Einen Schlüssel besitzt unser Betrieb nicht.",
         "Falls Sie zunächst nur eines der beiden Fahrzeuge oder ein kürzeres Zeitfenster prüfen können, bitten wir um gesonderte Rückmeldung. Einen allgemeinen Ersatzfuhrpark beantragen wir nicht. Die Fahrzeugdaten sind am 24. August nachgereicht worden. Den Lieferschein vom 19. August und eine getrennte Ladezeiterfassung für diese Woche stellen wir für den Ortstermin bereit.",
         "Mit freundlichen Grüßen\nHenning Kruse\nGeschäftsführer")
    save_doc(d, 11)

    sender = ("STADT EICHENSTEDT", "Tiefbauamt | Straßenbestandsakten", "Marktplatz 1, 38118 Eichenstedt", "Hauspost an die Straßenverkehrsbehörde")
    d = letter(15, sender, "Straßenverkehrsbehörde\nHerrn Timo Mäßner", "26. August 2026", "Straßenbestandsblatt BW-18 / zu 66.2-VA-2026-884", "Auszug aus der Straßenbestandsakte")
    body(d, "Auf Ihre Anfrage zum Buchenweg und Lindenplatz übersende ich die nachstehende Abschrift des Eintrags im Straßenbestandsblatt BW-18. Die Stadt Eichenstedt liegt im Land Niedersachsen. Die Bestandsakte wird nach dem Niedersächsischen Straßengesetz (NStrG) geführt.")
    heading(d, "1 Übertragener Widmungseintrag")
    body(d, "Der Buchenweg ist zwischen Lindenplatz und Einmündung Finkenstraße als Gemeindestraße für den öffentlichen Verkehr gewidmet. Die Widmung wird im Bestandsblatt auf Paragraf 6 NStrG gestützt. Als Straßenbaulastträgerin ist die Stadt Eichenstedt eingetragen. Der Eintrag erfasst Fahrbahn, beidseitige Gehwege und die straßenbegleitende Ladebucht vor Hausnummer 18. Eine Beschränkung auf einzelne Verkehrsarten ist im übertragenen Eintrag nicht vermerkt.",
         "Für den Lindenplatz ist die befestigte Randfläche an der Einmündung Buchenweg ebenfalls der Gemeindestraße zugeordnet. Die Busaufstellfläche ist als Verkehrsfläche verzeichnet. Der Eintrag enthält keine einem einzelnen Gewerbebetrieb zugewiesene Ladefläche und kein ausschließliches Nutzungsrecht der Bäckerei.")
    heading(d, "2 Ablage und Abgrenzung")
    body(d, "Die Abschrift gibt den am 26. August eingesehenen Bestandseintrag wieder. Das historische Bekanntmachungsblatt und der damalige Lageplan liegen in einem noch nicht beigezogenen Papierband des Stadtarchivs. Datum und Wortlaut der Bekanntmachung können aus dem Bestandsblatt allein nicht bestätigt werden. Eine beglaubigte Abschrift der vollständigen Widmungsverfügung ist dies nicht.",
         "Die Grenze zwischen Gehweg und privater Ladenschwelle an Hausnummer 18 ist aus der verfügbaren Übersicht nicht zentimetergenau bestimmbar. Für etwaige betriebliche Aufstellflächen wäre deshalb eine örtliche Abgrenzung erforderlich. Der Verkehrsversuch vom 12. August ist in einer gesonderten Verkehrsakte geführt. Die drei Poller am Gehwegzugang Lindenplatz stehen im Ausstattungsverzeichnis des städtischen Bauhofs; ein betrieblicher Schlüsselinhaber ist dort nicht eingetragen.",
         "Im Auftrag\nAnja Voss\nTiefbauamt")
    save_doc(d, 15)

    sender = ("STADT EICHENSTEDT", "Kommunale Verkehrsüberwachung", "Industrieweg 4, 38120 Eichenstedt", "Auszug für die Straßenverkehrsbehörde")
    d = letter(17, sender, "Straßenverkehrsbehörde\nVorgang Schulstraße Buchenweg", "26. August 2026", "OWi 26-441887 / Einsatz 19.08.2026", "Auszug aus dem Kontrollprotokoll")
    body(d, "Dieser Auszug überträgt die auf den Lieferwagen BS-KR 418 bezogenen Einträge aus dem Einsatzprotokoll vom 19. August. Die Uhrzeiten der beiden Bilddateien sind getrennt von den handschriftlich auf volle Minuten notierten Zeiten wiedergegeben.")
    heading(d, "1 Einträge am 19. August")
    body(d, "07:43:12 Uhr: BS-KR 418 wird aus Richtung Lindenplatz kommend vor Hausnummer 11 aufgenommen. Auf dem Bild ist eine männliche Person am Steuer zu erkennen. Ein Gespräch dieser Person mit der am Einmündungsbereich eingesetzten Mitarbeiterin ist im hier ausgewerteten Protokollabschnitt nicht dokumentiert.",
         "07:46:08 Uhr: Das Fahrzeug steht vor Hausnummer 18. Kisten werden zur Bäckerei getragen. Im erfassten Bereich der Windschutzscheibe ist kein Genehmigungsschreiben zu erkennen. Eine vollständige Durchsicht des Fahrzeuginnenraums erfolgte nicht.",
         "07:53 Uhr: Das Fahrzeug fährt in Richtung Finkenstraße ab. Im Notizfeld steht nach Ansprache des Fahrers: 'Sven Rüß, Firma Kruse, Antrag am Freitag geschickt'. Der Eintrag enthält keine Ausweiskopie und keine abschließende Feststellung zur Identität des Fahrers. Eine Behinderung durch ausweichende Kinder ist im Notizfeld nicht beschrieben.")
    heading(d, "2 Übertragene Unterlagen")
    body(d, "Die Bilddateien mit den Zeiten 07:43:12 und 07:46:08 Uhr sind in der Bildablage der Bußgeldstelle unter OWi 26-441887 verzeichnet. Sie wurden mit dieser Hauspost nicht erneut exportiert. Der übersandte Auszug ist ein Textauszug, kein Ersatz der beiden Übersichtsaufnahmen. Für deren erneute Übermittlung ist Frau Brüggemann angesprochen.",
         "Die Anhörung der Bußgeldstelle trägt das Datum 20. August 2026. Ein Zugang beim Betrieb ist aus diesem Einsatzprotokoll nicht ersichtlich. Die auf den Antrag bezogene Liste der Straßenverkehrsbehörde gehört nicht zum hiesigen Einsatzprotokoll.",
         "Übertragen und abgeglichen am 26. August 2026\nJana Brüggemann\nBußgeldstelle")
    save_doc(d, 17)

    d = letter(19, CITY, "Verkehrsakte Buchenweg\nKopie an Tiefbauamt und Bäckerei Kruse & Sohn KG", "27. August 2026", "66.2-VA-2026-884-A03", "Vermerk über den Ortstermin am Lindenplatz")
    body(d, "Der Ortstermin fand von 07:15 bis 08:18 Uhr statt. Anwesend waren Timo Mäßner, Anja Voss vom Tiefbauamt, Lea Drews von der Verkehrsüberwachung sowie Nele Maaß und zeitweise Petra Maaß für die Bäckerei. Miriam Eßwein kam um 07:35 Uhr hinzu. Die Filialleiterin und die Fahrerin sind zwei verschiedene Personen.")
    heading(d, "1 Fläche und Lieferweg")
    body(d, "Der Container stand bei Beginn des Termins unverändert an der Ersatzladefläche. Die beiden in früheren Nachrichten beschriebenen Pkw standen heute nicht dort. Ein Pkw hielt von 07:28 bis 07:34 Uhr im noch nutzbaren Rest der Markierung. Die Restlänge wurde längs des Bordes mit rund 4,90 Metern gemessen. BS-KR 422 ist nach den nachgereichten Fahrzeugdaten 5,41 Meter lang. Ein Entladen in der Bushaltestelle wurde nicht versucht.",
         "Frau Voss und Frau Nele Maaß gingen den Weg von der vorgesehenen Ladeposition bis zur Ladentür mit einem Messrad ab. Sie notierten 228 Meter einschließlich zweier Straßenquerungen. Der früher vom Betrieb genannte Wert von etwa 230 Metern ist damit als gerundete Weglänge nachvollziehbar. Die Messung erfolgte ohne Kartierung der Grundstücksgrenzen. Die engste Passage zwischen zwei Pollern am Gehwegzugang betrug 1,18 Meter; die Sackkarre ist etwa 0,62 Meter breit.")
    heading(d, "2 Beobachtete Übergabe")
    body(d, "BS-KR 422 wartete außerhalb der Bushaltestelle auf einem freien Stellplatz weiter am Lindenplatz. Petra Maaß begann dort um 07:44 Uhr mit dem Entladen. Die Ware wurde mit der Sackkarre in vier Gängen gebracht. Der letzte Eingang in der Filiale wurde um 08:07 Uhr notiert. Zwei Pausen von zusammen ungefähr drei Minuten entfielen auf den Busausstieg und eine Gruppe am Poller. Es kam zu keinem Zusammenstoß. Diese Zeiten enthalten den längeren Weg vom Ausweichstellplatz und sind keine Messung einer freien Ersatzladezone.",
         "Die Straßenverkehrsbehörde erteilte bei diesem Termin keine Zufahrtserlaubnis. Frau Voss prüft die Containerstellung und den Schlosszustand des mittleren Pollers. Frau Nele Maaß regte eine kleine verschließbare Rollbox an der Fassade an, damit die Verkäuferin nicht zwischen Kasse und Warenannahme wechseln muss. Ein Aufstellen wurde nicht vereinbart. Maße und Betriebszeiten sollen schriftlich vorgelegt werden.",
         "Anlage: Maßskizze vom 27. August 2026, ein Blatt.\nTimo Mäßner")
    save_doc(d, 19)

    d = letter(24, BAKERY, "Stadt Eichenstedt\nTiefbauamt und Straßenverkehrsbehörde\nMarktplatz 1\n38118 Eichenstedt", "31. August 2026", "Buchenweg 18 / Ergänzung zu 66.2-VA-2026-884-A03", "Antrag auf eine vorübergehende Rollboxstellfläche")
    body(d, "Sehr geehrte Damen und Herren,",
         "ergänzend zur beantragten Fahrzeugzufahrt beantragen wir eine befristete Stellfläche für eine verschließbare Rollbox an der Fassade unserer Filiale Buchenweg 18. Die Box soll vom 7. September bis 2. Oktober 2026 montags bis freitags jeweils von 07:25 bis 08:30 Uhr außen stehen. Sie soll auch dann noch dort bleiben, wenn das Lieferfahrzeug bereits abgefahren ist. Danach wird sie in die Filiale geschoben.",
         "Die Box hat eine Grundfläche von 1,20 mal 0,80 Metern, ist 1,45 Meter hoch und wiegt leer 58 Kilogramm. Sie soll mit ihrer langen Seite an der Wand rechts neben der Ladentür stehen. Nach der Messung vom 27. August würden auf dem 2,65 Meter breiten Gehweg rechnerisch 1,85 Meter verbleiben. Die Tür der Box öffnet nach oben. Weder Werbung noch ein Verkauf aus der Box sind vorgesehen. In ihr sollen nur geschlossene Brotkisten stehen; die Kühlboxen nehmen wir unmittelbar in den Laden.",
         "Die Räder können festgestellt werden. Wir werden keine Dübel, Fundamente oder sonstigen Befestigungen im Gehweg anbringen. Die Filialleiterin kontrolliert die Fläche vor dem Öffnen und räumt die Box bei einer erkennbaren Behinderung sofort hinein. Sie ist während der beantragten Zeit im Laden erreichbar. Die Schwelle und die Grenze zur öffentlichen Fläche sind uns nicht genau bekannt. Bitte beziehen Sie die beantragte Aufstellung deshalb vollständig in Ihre Prüfung ein.",
         "Die Box ist nur eine Möglichkeit für die Warenannahme. Eine Genehmigung für das Befahren der Schulstraße oder eine Öffnung der Poller setzen wir damit nicht voraus. Der Lieferweg vom Lindenplatz wäre mit der vorhandenen Sackkarre weiterhin möglich. Wir bitten auch um Mitteilung, ob für die Aufstellung ein weiterer Antrag oder ein Versicherungsnachweis benötigt wird. Eine konkrete Gebührenauskunft liegt uns noch nicht vor.",
         "Die Box ist bislang weder bestellt noch aufgestellt. Für die ersten zwei Wochen können wir täglich festhalten, wann sie außen stand und ob Beschwerden eingingen. Die Maße beziehen sich auf das von uns ausgewählte Angebot; bei einer anderen Ausführung würden wir vorher neue Maße einreichen.",
         "Mit freundlichen Grüßen\nHenning Kruse\nGeschäftsführer")
    save_doc(d, 24)

    d = letter(26, CITY, "Bäckerei Kruse & Sohn KG\nHerrn Henning Kruse\nMühlenstraße 7\n38118 Eichenstedt", "1. September 2026", "66.2-VA-2026-884-A03 / Flächenvorgang BW-18-R", "Stand Ihrer Anträge für den Buchenweg")
    body(d, "Sehr geehrter Herr Kruse,",
         "Ihre Ergänzung zur Fahrzeugzufahrt vom 26. August und Ihr Antrag zur Rollbox vom 31. August sind eingegangen. Die Vollständigkeitsmitteilung vom 25. August bezog sich auf den damaligen Zufahrtsantrag einschließlich der beiden Fahrzeugnachweise. Die neue Aufstellung einer Box wird zusätzlich unter dem Flächenvorgang BW-18-R erfasst und an das Tiefbauamt weitergegeben.",
         "Die Zufahrt wird weiterhin im Verfahren nach Paragraf 46 Absatz 1 Satz 1 Nummer 11 der Straßenverkehrs-Ordnung bearbeitet. Zum Flächenvorgang ist die Prüfung nach Paragraf 18 des Niedersächsischen Straßengesetzes (NStrG) aufgenommen worden. Welche weiteren Abstimmungen für die konkret beantragte Aufstellung erforderlich sind, ist noch nicht abschließend geklärt. Eine Genehmigung für eine der Nutzungen wird mit diesem Schreiben nicht erteilt.",
         "Bitte teilen Sie bis zum 4. September mit, ob Sie den Zufahrtsantrag bei Wegfall der morgendlichen Leergutrücknahme auf 07:35 bis 07:45 Uhr beschränken würden. Die Schule hat uns am 31. August ihre Beobachtungen übersandt. Ihr Ladezeitblatt enthält sowohl Fahrten vom Ausweichstellplatz als auch eine erst nach 08:15 Uhr erfolgte Direktanlieferung. Für einen Vergleich benötigen wir noch einen Morgen mit freier Ersatzladefläche.",
         "Zum Boxenstandort fehlen eine Bestätigung zur örtlichen Grenze an der Ladentür und die Benennung der Person, die bei Abwesenheit von Frau Maaß die Box räumen kann. Ob der rechnerisch verbleibende Gehweg im Betrieb tatsächlich frei bleibt, soll bei einem kurzen gemeinsamen Termin überprüft werden. Bitte stellen Sie die Box bis dahin nicht auf.",
         "Das Tiefbauamt hat die Versetzung des Containers für den 2. September angekündigt, jedoch noch nicht als ausgeführt gemeldet. Die Herausnahme des mittleren Pollers ist ebenfalls nicht freigegeben. Eine Schlüsselübergabe ist nicht vereinbart. Über den Bearbeitungsstand der übrigen sieben Anträge kann derzeit noch keine gemeinsame Entscheidung mitgeteilt werden.",
         "Eine abschließende Entscheidung liegt damit weiterhin nicht vor. Die angekündigte Entscheidungswoche läuft; einen verbindlichen Ausgabetermin kann ich Ihnen heute nicht nennen. Zur Bußgeldanhörung vom 20. August trifft dieses Schreiben keine Aussage.",
         "Mit freundlichen Grüßen\nIm Auftrag\nTimo Mäßner\nStadtoberinspektor")
    save_doc(d, 26)


def create_messages() -> None:
    email(16, "Bernd Semmler <b.semmler@kruse-brot.de>", "Henning Kruse <h.kruse@kruse-brot.de>", "Timo Mäßner <t.maessner@eichenstedt.de>", "Wed, 26 Aug 2026 08:42:16 +0200", "Buchenweg: Eintrag A-31 und unser Telefonat vom Freitag", """
Hallo Henning, sehr geehrter Herr Mäßner,

ich habe die Zeile zum 24. August in unserer zusammengeführten Lieferliste nachgesehen. A-31 ist bei uns die interne Tourreferenz für die Montagstour nach dem Fahrzeugtausch. Ich habe sie beim Zusammenführen in die Spalte Genehmigung übernommen und daneben 'erstmals genehmigt' gesetzt, weil ich unsere interne Freigabe der Tour mit einer städtischen Freigabe verwechselt habe. Das war mein Fehler. Einen Bescheid A-31 kann ich nicht vorlegen. In unserem Postfach liegt nur die Zwischennachricht der Stadt vom 25. August.

Die Zeiten 07:39 und 07:49 stammen aus Svens Tourzettel. Ich ändere die versandte Liste nicht rückwirkend. Der Genehmigungseintrag und die Bemerkung in dieser Zeile sind nicht als Nachweis einer Erlaubnis verwendbar. Sven ist an diesem Montag tatsächlich gefahren. Eine zusätzliche Anlieferung am Lindenplatz an diesem Tag ist damit nicht gemeint. Ich habe den beiden Fahrern heute gesagt, dass bis zu einer schriftlichen Entscheidung keine weitere Einfahrt während der Sperrzeiten eingeplant wird.

Henning hat mir außerdem gesagt, er habe am Freitag am Telefon von einer 'Anhörung am Mittwoch' gesprochen. Gemeint war nach seiner heutigen Erinnerung, dass Sven am Mittwoch vor Ort angesprochen und aufgeschrieben wurde. Das Schriftstück ist auf den 20. August datiert. Wer es wann aus dem Briefkasten genommen hat, können wir nicht mehr sicher sagen. Ich kann daher nicht bestätigen, dass am Mittwoch bereits die schriftliche Anhörung vorlag.

Den Lieferschein 2026-08-19-031 habe ich im Warensystem gefunden und als Nachdruck zur Ablage gegeben. Svens Foto von 07:40 Uhr ist im gemeinsamen Ordner nicht vorhanden; er sucht es noch auf seinem alten Telefon. Die am 14. August versandten Dateien Antrag_Ausnahme_Buchenweg.pdf und Lieferplan_Filiale_Buchenweg.xlsx liegen in meinem jetzigen Export ebenfalls nicht. Die zusammengeführte Lieferliste ist nicht identisch mit dieser ursprünglichen Anlage. Ich werde fehlende Originalanhänge nicht durch eine neu erstellte Datei gleichen Namens ersetzen.

Viele Grüße
Bernd Semmler
Disposition / Fuhrpark
Bäckerei Kruse & Sohn KG
Mühlenstraße 7, 38118 Eichenstedt
Telefon 0531 882 121
""", "<77c2f1b09aa3e5d21c04@kruse-brot.de>")

    email(18, "Timo Mäßner <t.maessner@eichenstedt.de>", "Henning Kruse <h.kruse@kruse-brot.de>", "Anja Voss <a.voss@eichenstedt.de>", "Wed, 26 Aug 2026 14:18:09 +0200", "Ortstermin 27. August: Lieferablauf und Unterlagen", """
Sehr geehrter Herr Kruse,

für morgen bleibt es bei 07:15 Uhr am Lindenplatz. Ihre Fahrzeugunterlagen sind vollständig, wie am 25. August mitgeteilt. Die folgenden Fragen betreffen den tatsächlichen Lieferablauf und ändern diese Eingangsbestätigung nicht.

Bitte lassen Sie uns die tatsächlich benötigte Standzeit, die Zahl der Transportgänge und eine mögliche Trennung von Frischware und Leergut nennen. Ein pauschales Zeitfenster von zwanzig Minuten sagt noch nicht, wann das Fahrzeug vor der Tür steht. Interessant ist auch, ob die Verkäuferin die Annahme allein erledigt oder dafür den Laden verlassen muss. Eine Fahrt in der Sperrzeit zum Zweck der Vorführung ist nicht vereinbart.

Das Tiefbauamt bringt ein Messrad mit. Bitte halten Sie Ihre Sackkarre bereit, damit die Durchfahrt zwischen den Pollern und der Weg bis zum Laden nachvollzogen werden können. Die Skizze VA-884 Blatt 2 liegt nun als Planabschrift vor. Sie enthält keine technischen Maße zur Containerstellung; dafür wird beim Termin ein eigenes Blatt angelegt.

Die Berichtigung von Herrn Semmler habe ich zur Akte genommen. Unter A-31 liegt in unserer Antragsakte kein für Ihren Betrieb erteilter Bescheid. Die Apotheke ist ein anderer Vorgang; aus dem dort sichtbaren Schreiben kann ich keine Freigabe für Ihre Fahrzeuge ableiten.

Noch nicht in der hier zusammengestellten Arbeitsablage liegen der vollständige Anliegerverteiler und das leere Formblatt aus dem Schreiben vom 12. August. Diese habe ich intern angefordert. Frau Eßwein hat mitgeteilt, dass ihr ursprüngliches Fotoarchiv und die einzelne Zählliste vom 18. August erneut vom Elternratsrechner exportiert werden müssen. Der Rohdatenexport mit mehreren Zähltagen ersetzt die ursprüngliche Zählliste nicht.

Mit freundlichen Grüßen
Timo Mäßner
Stadtoberinspektor
Stadt Eichenstedt, Straßenverkehrsbehörde
Marktplatz 1, 38118 Eichenstedt
Telefon 0531 771 362
""", "<buchenweg-16-202608@kruse-brot.de>")

    email(22, "Anja Voss <a.voss@eichenstedt.de>", "Timo Mäßner <t.maessner@eichenstedt.de>", "Henning Kruse <h.kruse@kruse-brot.de>", "Fri, 28 Aug 2026 11:36:40 +0200", "Lindenplatz: Containertermin und Pollerschloss", """
Guten Tag Herr Mäßner,

der Bauleiter hat heute bestätigt, dass der Container noch bis Freitag, 4. September, für die Kanalsanierung gebraucht wird. Die zunächst besprochenen sechs Meter Verschiebung Richtung Finkenstraße würden den Zugang zum Schacht zustellen. Nach dem Ortstermin ist stattdessen eine seitliche Umsetzung innerhalb der Baustellenfläche möglich. Dafür muss zunächst das Materiallager umgeräumt werden.

Der Unternehmer hat den Lkw für Mittwoch, 2. September, zwischen 10 und 12 Uhr angekündigt. Wenn die Umsetzung wie besprochen gelingt, soll die markierte Ladefläche wieder auf ihrer ganzen Länge von acht Metern frei sein. Das ist eine Terminankündigung, keine Vollzugsmeldung. Für Montag und Dienstag kann ich die Freigabe nicht zusagen. An den beiden Tagen bleibt die Lage vom Ortstermin maßgeblich.

Der mittlere der drei Poller am Gehwegzugang hat ein altes Dreikantschloss. Der Bauhof konnte ihn heute nicht ziehen; die Bodenhülse sitzt fest. Ein Schlüssel allein löst das Problem nicht. Eine Öffnung für die Bäckerei ist weder erprobt noch angeordnet. Die Poller sperren nicht die Fahrbahn des Buchenwegs. Eine aufgeräumte Passage für Fußgänger und Sackkarren darf nicht als neue Kfz-Zufahrt verstanden werden.

Die von Frau Maaß erwähnte Rollbox an Hausnummer 18 ist auf unserer Maßskizze nur als besprochene Fläche dargestellt. Die Gehwegbreite von 2,65 Metern wurde ab Fassade gemessen, nicht anhand einer Katastergrenze. Eine Zusage zur Aufstellung habe ich nicht gegeben. Bitte lassen Sie uns die genaue Nutzungsdauer und Abmessungen schriftlich zukommen.

Viele Grüße
Anja Voss
Tiefbauamt, Stadt Eichenstedt
Marktplatz 1, 38118 Eichenstedt
Hausdurchwahl 417
""", "<buchenweg-18-202608@eichenstedt.de>")

    email(25, "Miriam Eßwein <m.esswein@elternrat-buchenweg.de>", "Timo Mäßner <t.maessner@eichenstedt.de>", "schulleitung@gs-buchenweg.de, Henning Kruse <h.kruse@kruse-brot.de>", "Mon, 31 Aug 2026 09:24:33 +0200", "Lieferfenster Buchenweg: Rückmeldung nach dem Ortstermin", """
Guten Morgen Herr Mäßner,

ich habe nach dem Termin mit der Schulleitung und den beiden Eltern gesprochen, die morgens am Lindenplatz stehen. Wir möchten, dass die Filiale beliefert werden kann. Der jetzige Weg mit der Sackkarre ist aber gerade beim Busausstieg unübersichtlich. Dass es beim Termin keinen Zusammenstoß gab, heißt nicht, dass vier Gänge mit Ware morgens problemlos sind.

Die Buslinie 16 kommt an den von uns beobachteten Schultagen meist zwischen 07:37 und 07:43 Uhr an. Das sind unsere Beobachtungen und keine Fahrplanauskunft. Auf dem Gehweg warten manche Kinder nach dem Aussteigen auf Geschwister. Eine längere Öffnung zwischen den Pollern würde das Gedränge verringern, wir möchten dort aber keine Fahrzeuge und auch keine unbeaufsichtigt aufgestellten Kisten.

Ein festes Lieferfenster könnte verständlicher sein als die täglichen Ausnahmen, über die wir bisher nur gerüchteweise hören. Ob 07:35 bis 07:45 Uhr weniger Konflikte macht, können wir ohne Beobachtung nicht sagen: Es überlappt gerade mit dem Bus. Nach 08:15 Uhr ist es erheblich ruhiger, allerdings verstehen wir, dass die Backwaren dann spät ankommen. Von einer generellen Zustimmung des Elternrats zu einem bestimmten Fenster bitte ich deshalb abzusehen.

Vor dem Laden führt die Schlange der wartenden Kunden manchmal an der Fassade entlang. Wenn dort zusätzlich eine Box steht, interessiert uns die tatsächlich verbleibende Passage. Bitte messen Sie nicht nur die freie Gehwegbreite bei geschlossenem Laden. Die Namen einzelner Kinder sollen weiterhin nicht in die Unterlagen aufgenommen werden.

Mit freundlichen Grüßen
Miriam Eßwein
Vorsitzende des Schulelternrats
Grundschule Am Buchenweg
""", "<2abbddadeaf18c274d47@elternrat-buchenweg.de>")


def create_plain_records() -> None:
    (CASE / FILES[12]).write_text("""STADT EICHENSTEDT
Straßenverkehrsbehörde | Eingangsjournal 66.2-VA-2026-884-A03
Auszug zum Nachreichungseingang am 24. August 2026, 08:16 Uhr
Bearbeiter: Timo Mäßner. Übertragung in die Arbeitsablage: 26. August 2026.

Der Betrieb Kruse & Sohn KG hat die Kopie der Zulassungsbescheinigung Teil I
für BS-KR 422 nachgereicht. Der am 14. August versandte Sammelscan enthielt
nach Durchsicht zwei Seiten zum Fahrzeug BS-KR 418 und keine Seite zu
BS-KR 422. Der Austausch wurde nur in der Fahrzeugablage vorgenommen;
der ursprüngliche E-Mail-Eingang bleibt unverändert gespeichert.

BS-KR 418: Halter Bäckerei Kruse & Sohn KG, Mühlenstraße 7,
38118 Eichenstedt. Kastenwagen, zulässige Gesamtmasse 3.500 kg.
Länge laut Fahrzeugnachweis 5.998 mm, Breite ohne Spiegel 2.050 mm.
Die Disposition nennt eine Breite mit Spiegeln von ungefähr 2.430 mm.

BS-KR 422: gleicher Halter. Kastenwagen, zulässige Gesamtmasse 3.500 kg.
Länge laut Fahrzeugnachweis 5.413 mm, Breite ohne Spiegel 2.050 mm.
Die Disposition nennt eine Breite mit Spiegeln von ungefähr 2.470 mm.
Die Spiegelmaße sind Betriebsangaben und nicht dem amtlichen Feld für
die Fahrzeugbreite entnommen.

Die wechselnden Fahrer sind Sven Rüß und Petra Maaß. Eine gleichzeitige
Zufahrt beider Fahrzeuge ist nach Rückfrage bei der Disposition nicht
beabsichtigt. Für Ersatzfahrzeuge ohne eines dieser beiden Kennzeichen
ist bislang kein Nachweis eingereicht.

Der neue Fahrzeugschein ist lesbar. Die Vollständigkeitsmitteilung zum
bisherigen Zufahrtsantrag wird vorbereitet. Eine Entscheidung über die
Zufahrt ist mit dieser Erfassung nicht verbunden.

Ablagehinweis: Der ursprüngliche Sammelscan und die nachgereichte Kopie
verbleiben im Register Fahrzeugnachweise der Verkehrsakte. Sie wurden
bei diesem Textauszug nicht mit exportiert. Dieser Auszug enthält die
übertragenen Fahrzeugdaten, nicht die Bildwiedergabe der Dokumente.

Mäßner
""", encoding="utf-8")

    (CASE / FILES[23]).write_text("""Bäckerei Kruse & Sohn KG | Filiale Buchenweg 18
Notiz von Nele Maaß an Henning Kruse und Bernd Semmler
Freitag, 28. August 2026, 13:10 Uhr

Ich habe die Ladezeiten vom 25. bis 28. August nach unseren Aufzeichnungen
in eine Tabelle übertragen. Start ist jeweils das Öffnen der Ladetür am
stehenden Fahrzeug. Ende ist das Abstellen der letzten Kiste im Laden.
Die Zahlen sind Minuten nach unserer Wanduhr beziehungsweise Petras Uhr;
wir haben die Uhren nicht synchronisiert. Stand- und Wegezeiten vor dem
ersten Entladen gehören nicht dazu. Die Behörde hat am Donnerstag selbst
mitgeschrieben, die übrigen Tage stammen von uns.

Am Dienstag kam Sven am Lindenplatz auf einem Ausweichstellplatz zum
Stehen. Ich konnte wegen der Kunden nicht mitgehen. Er brauchte vier
Gänge. In der Kasse sind zwischen 07:42 und 07:59 Uhr zwölf Verkäufe
gebucht. Eine Kundin nahm Toast statt des noch fehlenden Brots. Eine
bezifferte Umsatzeinbuße kann ich daraus nicht ableiten.

Am Mittwoch brachten wir haltbare Ware um 06:58 Uhr direkt vor den Laden.
Die Frischware kam getrennt vom Lindenplatz. Damit war die Auslage früher
teilweise gefüllt, aber der zweite Weg entfiel nicht. Bernd hatte die
leeren Kisten schon am Vortag abholen lassen. Die Kühlboxen kamen zuerst
rein und wurden sofort in den Kühlraum gebracht. Temperaturen habe ich
in diesem Ladezeitblatt nicht erfasst; aus den Minuten allein folgt
kein Nachweis einer unterbrochenen Kühlkette.

Am Donnerstag während des Ortstermins war Petra um 08:07 Uhr mit allem
fertig. Bei meinem Wechsel zwischen Tür und Kasse blieb die Ladentür
zweimal länger offen. Die Kindergruppe am Poller habe ich nicht vom Laden
aus gesehen; die Angabe dazu stammt von Petra und vom Ortstermin.

Heute haben wir die Frischware erst nach Ende der morgendlichen Sperrzeit
direkt anliefern lassen. Sven entlud von 08:18 bis 08:25 Uhr. Das ging
schneller, aber bis dahin fehlten die Nordring-Brote. Vier Stammkunden
fragten danach. Bestellte belegte Brötchen waren aus der frühen Lieferung
vorhanden. Das war kein vollständiger Ausfall des Frühstücksangebots.

Die geplante Rollbox soll den Wechsel zwischen Kundenbedienung und
Warenannahme erleichtern. Sie verkürzt nicht den Weg vom Lindenplatz.
Ich kann sie vor dem Laden beaufsichtigen, solange ich Dienst habe.
Am Mittwoch beginnt mein Dienst erst um 08:00 Uhr; für diese Zeit müsste
noch jemand benannt werden. Wir haben noch keine Box im Laden.

Nele Maaß
""", encoding="utf-8")
    rows = [
        ["2026-08-25", "BS-KR 418", "Lindenplatz Ausweichstellplatz", "07:39", "08:03", "4", "16 Brotkisten und 2 Kühlboxen", "Nele Maaß / Sven Rüß; erster und letzter Wareneingang"],
        ["2026-08-26", "BS-KR 422", "Buchenweg 18 direkt", "06:58", "07:06", "2", "8 Kisten haltbare Ware", "Nele Maaß / Petra Maaß; vor Sperrzeit"],
        ["2026-08-26", "BS-KR 422", "Lindenplatz Ausweichstellplatz", "07:42", "08:01", "3", "10 Brotkisten und 2 Kühlboxen", "Nele Maaß / Petra Maaß; keine Leergutrücknahme"],
        ["2026-08-27", "BS-KR 422", "Lindenplatz Ausweichstellplatz", "07:44", "08:07", "4", "14 Brotkisten und 2 Kühlboxen", "Ortstermin; ungefähr 3 Minuten Querungswartezeit enthalten"],
        ["2026-08-28", "BS-KR 418", "Buchenweg 18 direkt", "06:57", "07:04", "2", "8 Kisten haltbare Ware", "Nele Maaß / Sven Rüß; vor Sperrzeit"],
        ["2026-08-28", "BS-KR 418", "Buchenweg 18 direkt", "08:18", "08:25", "2", "12 Brotkisten und 2 Kühlboxen", "Nele Maaß / Sven Rüß; nach Sperrzeit"],
    ]
    with (CASE / FILES[21]).open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=";", lineterminator="\r\n")
        writer.writerow(["Datum", "Fahrzeug", "Entladeort", "Beginn_Entladung", "Letzter_Wareneingang", "Transportgänge", "Ladung", "Quelle_und_Abgrenzung"])
        writer.writerows(rows)


PDF_BODY = ParagraphStyle("body", fontName="Times-Roman", fontSize=11, leading=14)


def register_pdf_fonts() -> None:
    font_dir = Path("/System/Library/Fonts/Supplemental")
    for alias, name in (("Times-Roman", "Times New Roman.ttf"), ("Times-Bold", "Times New Roman Bold.ttf")):
        font = font_dir / name
        if not font.is_file():
            raise RuntimeError(f"Schriftdatei für eingebettete PDF-Schrift fehlt: {font}")
        pdfmetrics.registerFont(TTFont(alias, str(font)))


def pdf_start(number: int, title: str, subtitle: str) -> canvas.Canvas:
    c = canvas.Canvas(str(CASE / FILES[number]), pagesize=A4, invariant=1)
    c.setTitle(title)
    c.setAuthor("Stadt Eichenstedt" if number != 13 else "Bäckerei Kruse & Sohn KG")
    c.setFont("Times-Bold", 16); c.drawString(48, 792, title)
    c.setFont("Times-Roman", 11); c.drawString(48, 770, subtitle)
    return c


def pdf_paragraph(c: canvas.Canvas, text: str, y: float, width: float = 499) -> float:
    p = Paragraph(escape(text), PDF_BODY)
    _, h = p.wrap(width, 700)
    if y - h < 43:
        raise ValueError("PDF-Text überschreitet den Satzspiegel")
    p.drawOn(c, 48, y-h)
    return y-h-10


def pdf_end(c: canvas.Canvas, reference: str) -> None:
    c.setFont("Times-Roman", 9); c.drawString(48, 28, reference + " | Blatt 1")
    c.showPage(); c.save()


def create_delivery_note() -> None:
    c = pdf_start(13, "BÄCKEREI KRUSE & SOHN KG", "Mühlenstraße 7, 38118 Eichenstedt | Telefon 0531 882 117")
    c.setFont("Times-Bold", 14); c.drawString(48, 727, "Lieferschein 2026-08-19-031")
    y = pdf_paragraph(c, "Lieferung am 19. August 2026 | Tour 3 | Fahrzeug BS-KR 418", 704)
    y = pdf_paragraph(c, "Empfänger: Filiale Buchenweg 18, 38118 Eichenstedt. Fahrer: Sven Rüß. Warenannahme: Nele Maaß.", y)
    data = [["Ware", "Menge", "Transportbehälter"], ["Mischbrote 750 g", "48 Stück", "6 Brotkisten"], ["Brötchensortiment", "240 Stück", "6 Brotkisten"], ["Süßgebäck", "48 Stück", "2 Brotkisten"], ["Belegte Brötchen", "36 Stück", "2 Kühlboxen"]]
    table = Table(data, colWidths=[233, 96, 170], rowHeights=[27]*5)
    table.setStyle(TableStyle([("FONT", (0,0), (-1,-1), "Times-Roman", 11), ("FONT", (0,0), (-1,0), "Times-Bold", 11), ("BACKGROUND", (0,0),(-1,0), colors.HexColor("#eeeeee")), ("LINEBELOW", (0,0),(-1,0), .5, colors.black), ("LINEBELOW", (0,1),(-1,-1), .25, colors.grey), ("VALIGN", (0,0),(-1,-1),"MIDDLE")]))
    table.wrapOn(c,499,400); table.drawOn(c,48,y-135); y -= 158
    y = pdf_paragraph(c, "Behältersumme: 14 Brotkisten und zwei Kühlboxen. Leergut wurde bei dieser Übergabe nicht zurückgenommen. Die Kühlboxen wurden ungeöffnet an die Filialleiterin übergeben.", y)
    y = pdf_paragraph(c, "Im Warensystem erfasster Übergabevermerk: Fahrzeug vor dem Laden ab etwa 07:46 Uhr; letzte Kiste um etwa 07:52 Uhr übernommen. Die Minutenangaben wurden von der Filialleiterin nach der Annahme eingetragen. Sie sind keine elektronischen Standortdaten.", y)
    y = pdf_paragraph(c, "Annahmevermerk: Menge nach Behältern vollständig; einzelne Stückzahlen beim Auspacken geprüft. Es wurde bei der Übergabe keine Fehlmenge gemeldet. Die Ware wurde der Filiale zugeordnet, ein Barverkauf erfolgte nicht.", y)
    y = pdf_paragraph(c, "Nachdruck aus dem Warensystem vom 26. August 2026, 08:31 Uhr, veranlasst durch Bernd Semmler. Der Nachdruck enthält den gespeicherten Annahmevermerk, aber keine eingescannte Unterschrift. Er gibt keine Auskunft über eine Zufahrtserlaubnis.", y)
    pdf_end(c, "Tour 3 | Lieferung 19.08.2026 | Nachdruck 26.08.2026")


def label(c: canvas.Canvas, x: float, y: float, text: str, size: float = 10) -> None:
    c.setFont("Times-Roman", size)
    c.drawString(x, y, text)


def create_sign_plan() -> None:
    c = pdf_start(14, "Verkehrszeichenplan VA-884 Blatt 2", "Stadt Eichenstedt | Schulstraße Buchenweg | Anordnung vom 12. August 2026")
    y = pdf_paragraph(c, "Planabschrift vom 26. August 2026 aus der Verkehrsakte. Die Lage ist schematisch dargestellt, nicht maßstäblich. Die Skizze gibt die Zuordnung der Zufahrten und Lieferstellen wieder; sie ist kein vermessener Lageplan.", 744)
    c.setFillColor(colors.HexColor("#ededed")); c.rect(242,365,84,305,fill=1,stroke=0)
    c.rect(65,365,460,43,fill=1,stroke=0); c.rect(65,627,460,43,fill=1,stroke=0)
    c.setFillColor(colors.black)
    c.setLineWidth(.7); c.line(242,408,242,627); c.line(326,408,326,627)
    label(c,80,382,"Lindenplatz",12); label(c,80,644,"Finkenstraße",12)
    c.saveState(); c.translate(275,479); c.rotate(90); label(c,0,0,"Buchenweg",12); c.restoreState()
    c.rect(349,487,146,32,stroke=1); label(c,359,499,"Bäckerei, Hausnummer 18")
    c.rect(349,544,146,32,stroke=1); label(c,359,556,"Apotheke, Hausnummer 22")
    c.rect(69,492,143,45,stroke=1); label(c,78,518,"Grundschule"); label(c,78,504,"Am Buchenweg")
    c.rect(75,420,137,35,stroke=1); label(c,84,442,"Ersatzladefläche"); label(c,84,428,"am Lindenplatz")
    for x,yy,txt in [(232,418,"1"),(336,618,"2")]:
        c.circle(x,yy,9,stroke=1); label(c,x-2.8,yy-3,txt)
    c.setDash(4,3); c.line(302,423,302,609); c.setDash()
    c.line(302,609,298,600); c.line(302,609,306,600)
    label(c,345,466,"Ladeposition vor Nr. 18")
    label(c,345,451,"während Sperrzeit verlegt")
    y = 339
    for text in [
        "1 und 2 bezeichnen die Zugänge des geregelten Abschnitts. Gestrichelte Linie: vom Betrieb beschriebener Lieferweg mit Ausfahrt Finkenstraße. Die Linie ist keine Genehmigung und keine angeordnete Einbahnregelung.",
        "Zeittext laut Anordnung: An Schultagen werktags 07:15 bis 08:15 Uhr und 12:45 bis 14:15 Uhr. Ausgenommen sind Linienverkehr, Rettungsdienst, Fahrräder und Fahrzeuge mit erteilter Ausnahmegenehmigung. Versuchsdauer: 17. August bis 18. Dezember 2026.",
        "Die genaue Nummerierung und grafische Ausführung der Haupt- und Zusatzzeichen lässt sich aus der verfügbaren Plankopie nicht sicher übertragen. Sie wird deshalb hier nicht ergänzt. Die Aufstellliste des Bauhofs mit diesen Angaben ist angefordert. Aufstellung und Verhüllung waren für den 14. August vorgesehen.",
        "Container, Bushaltestelle und Poller sind auf diesem Planblatt nicht vermessen. Eine gesonderte Maßaufnahme ist für den Ortstermin am 27. August vorgesehen. Abschrift erstellt: Timo Mäßner, Straßenverkehrsbehörde.",
    ]: y = pdf_paragraph(c,text,y)
    pdf_end(c,"66.2-VA-2026-884 | Planabschrift 26.08.2026")


def create_measurement_sketch() -> None:
    c = pdf_start(20, "Lindenplatz und Ladenzugang", "Maßskizze zum Ortstermin am 27. August 2026 | Anja Voss / Nele Maaß")
    y = pdf_paragraph(c, "Drei örtliche Ausschnitte, nicht maßstäblich und nicht lagegetreu zueinander. Längen am Bord und Poller mit Maßband, Weglänge mit Messrad ermittelt. Es handelt sich nicht um eine Kataster- oder Ausführungszeichnung.", 744)
    c.setFont("Times-Bold",11); c.drawString(48,672,"1 Ersatzladefläche am Lindenplatz")
    c.rect(65,610,440,35,stroke=1)
    c.setFillColor(colors.HexColor("#dddddd")); c.rect(65,610,170,35,fill=1,stroke=1); c.setFillColor(colors.black)
    label(c,78,623,"Containerüberstand"); label(c,265,623,"Restlänge etwa 4,90 m")
    label(c,65,588,"Gesamte markierte Länge 8,00 m; Fläche am 27.08. nicht durchgehend frei.")
    label(c,65,572,"Bushaltestelle angrenzend, nicht als Ersatzposition vermessen oder freigegeben.")
    c.setFont("Times-Bold",11); c.drawString(48,537,"2 Poller am Gehwegzugang")
    for x in [105,276,451]: c.circle(x,496,8,stroke=1,fill=0)
    c.line(114,496,267,496); c.line(285,496,442,496)
    label(c,155,507,"1,18 m licht"); label(c,326,507,"1,21 m licht")
    label(c,70,469,"Fester Poller"); label(c,225,469,"Mittlerer Poller"); label(c,417,469,"Fester Poller")
    label(c,70,450,"Mittlerer Poller mit Schloss; am Ortstermin nicht herausgenommen.")
    label(c,70,434,"Die Passage liegt im Gehweg, nicht in der Fahrbahn des Buchenwegs.")
    c.setFont("Times-Bold",11); c.drawString(48,395,"3 Besprochener Boxenstandort vor Buchenweg 18")
    c.line(65,364,510,364); label(c,72,371,"Fassade; Grundstücksgrenze nicht festgestellt")
    c.rect(105,316,160,48,stroke=1); label(c,116,334,"Box 1,20 x 0,80 m")
    c.setDash(3,3); c.line(65,233,510,233); c.setDash()
    label(c,310,326,"Gehweg gesamt: 2,65 m")
    label(c,310,300,"Verbleibend: 1,85 m")
    label(c,72,244,"Bord zur Fahrbahn; Fahrbahnengstelle im Abschnitt laut Anordnung 5,10 m")
    y = 209
    for text in [
        "Der eingezeichnete Boxenumriss bezeichnet nur den beim Ortstermin besprochenen Standort. Am 27. August stand dort keine Box. Der formelle Antrag und die Nutzungszeiten lagen bei der Messung noch nicht vor.",
        "Weglänge von der vorgesehenen Ersatzladeposition bis zur Ladentür: 228 m mit zwei Straßenquerungen. Der am Ortstermin tatsächlich benutzte Ausweichstellplatz liegt weiter entfernt; dessen Weglänge wurde nicht gemessen.",
        "Alle Werte sind örtliche Einzelmessungen. Wartende Kunden, geöffnete Türen und querende Kinder wurden nicht als feste Flächen abgezogen. Eine Freigabe von Flächen oder Pollern ist mit der Skizze nicht verbunden.",
    ]: y = pdf_paragraph(c,text,y)
    pdf_end(c,"66.2-VA-2026-884-A03 | Ortstermin 27.08.2026")


def validate_new() -> None:
    for number, name in FILES.items():
        path = CASE / name
        if not path.is_file() or path.stat().st_size < 400:
            raise ValueError(f"Fehlendes oder leeres Aktenstück {name}")
        if path.suffix == ".docx":
            d = Document(path)
            text = "\n".join(p.text for p in d.paragraphs)
            assert len(text) > 1800, name
            assert all(abs(s.page_width.cm-21) < .02 and abs(s.page_height.cm-29.7) < .02 for s in d.sections), name
            assert not d.styles.element.xpath(".//w:pBdr"), name
        elif path.suffix == ".pdf":
            r = PdfReader(path)
            assert len(r.pages) == 1, name
            text = "\n".join(p.extract_text() for p in r.pages)
            assert len(text) > 600, name
        elif path.suffix == ".eml":
            msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            for key in ("From", "To", "Cc", "Date", "Subject", "Message-ID", "Reply-To", "MIME-Version", "Content-Type", "Content-Transfer-Encoding", "In-Reply-To", "References"):
                assert msg[key], (name, key)
            assert not msg.defects, (name, msg.defects)
            text = msg.get_content()
        else:
            text = path.read_text(encoding="utf-8")
        for forbidden in ("§", "Musterlösung", "Arbeitsauftrag", "KI generiert", "Testakte", "\ufffd"):
            assert forbidden not in text, (name, forbidden)
    with (CASE / FILES[21]).open(newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert len(rows) == 7 and all(len(r) == 8 for r in rows)
    for row in rows[1:]:
        datetime.strptime(row[0], "%Y-%m-%d")
        assert datetime.strptime(row[3], "%H:%M") < datetime.strptime(row[4], "%H:%M")
        assert int(row[5]) > 0


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--upload-zip", type=Path)
    args = parser.parse_args()
    verify_originals(args.upload_zip)
    if not args.check_only:
        register_pdf_fonts()
        create_letters()
        create_messages()
        create_plain_records()
        create_delivery_note()
        create_sign_plan()
        create_measurement_sketch()
    validate_new()
    verify_originals(args.upload_zip)
    print("10 Originale bytegleich; 16 fallbezogene Ergänzungen geprüft.")
    if args.check_only:
        print("Gelesener Ausgangsbestand: 5 DOCX, 1 XLSX mit 2 Blättern, 3 EML, 1 CSV.")
        print("Upload und ursprüngliche Repo-Gesamt-PDF: je 17 Seiten, unterschiedliche Bytes.")
        print("Neue Stücke: 6 DOCX, 3 PDF, 4 EML, 2 TXT, 1 CSV; Original-XLSX unverändert.")
    print("Layoutprüfung der DOCX- und PDF-Seiten ist zusätzlich erforderlich.")


if __name__ == "__main__":
    main()
