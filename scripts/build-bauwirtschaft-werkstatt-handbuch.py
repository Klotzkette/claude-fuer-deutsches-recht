#!/usr/bin/env python3
"""Baut die handkuratierte Bauwerkstatt aus 100 progressiv lesbaren Stationen.

Die kanonischen Texte stehen in bauwirtschaft/references/werkstatt. Ein enger
Skill lädt nur das passende Modul. Der eigenständige Markdown-Download und
das formatierte Handbuch enthalten sämtliche Stationen in derselben Fassung.
Keine Textgenerierung, keine Seitenauffüllung und keine künstliche Verkleinerung.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PLUGIN = ROOT / "bauwirtschaft"
SOURCE = PLUGIN / "references/werkstatt"
OUT = PLUGIN / "materialien"
STEM = "bauwirtschaft-werkstatt-100-seiten"
TITLE = "Bauwirtschaft Werkstatt vom Projektauftrag bis zum Abschluss"
INTRO = (
    "Autor: Klotzkette. Quellen- und Redaktionsstand: 28.09.2026. "
    "Hundert zusammenhängende Arbeitsstationen; nur die zum Auftrag passenden Wege bearbeiten. "
    "LPH 1 bis 9: Stationen 7 bis 60. Fachvertiefungen: 61 bis 94. "
    "Quellen und Abschluss: 95 bis 100. Hauptfall Vermietung; Verkauf nur als gesonderte Option."
)


def sources():
    result = []
    for path in sorted(SOURCE.glob("*.md")):
        text = path.read_text(encoding="utf-8")
        for match in re.finditer(r"^## (\d+)\. ([^\n]+)\n(.*?)(?=^## |\Z)", text, re.M | re.S):
            result.append((int(match[1]), match[2], match[3].strip(), path))
    if [item[0] for item in result] != list(range(1, 101)):
        raise ValueError("Genau 100 eindeutige Stationen in der Reihenfolge 1 bis 100 erforderlich")
    return result


def source_digest():
    h = hashlib.sha256()
    for path in sorted(SOURCE.glob("*.md")):
        h.update(path.name.encode()); h.update(b"\0"); h.update(path.read_bytes())
    return h.hexdigest()


def markdown(items):
    body = ["# 1. " + TITLE, "", INTRO, ""]
    for number, title, content, _ in items:
        body += [f"## {number}. {title}", "", content, ""]
    path = PLUGIN / "bauwirtschaft-werkstatt.md"
    path.write_text("\n".join(body).rstrip() + "\n", encoding="utf-8")
    return path


def add_text(paragraph, text):
    """Native klickbare Quellenlinks; lange URLs stehen nicht als Satzspiegeltext."""
    previous = 0
    for match in re.finditer(r"\[([^]]+)\]\((https?://[^)]+)\)", text):
        paragraph.add_run(text[previous:match.start()])
        relation = paragraph.part.relate_to(
            match[2], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
        )
        link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), relation)
        run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Times New Roman"); fonts.set(qn("w:hAnsi"), "Times New Roman")
        props.append(fonts)
        color = OxmlElement("w:color"); color.set(qn("w:val"), "17365D"); props.append(color)
        underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); props.append(underline)
        run.append(props); txt = OxmlElement("w:t"); txt.text = match[1]; run.append(txt)
        link.append(run); paragraph._p.append(link); previous = match.end()
    paragraph.add_run(text[previous:])


def word(items):
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document(); section = doc.sections[0]
    # A4 entspricht der bestehenden deutschen Bauwerkstatt und ihren Vorlagen.
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(1.75); section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.9); section.right_margin = Cm(1.9)
    section.header_distance = Cm(.65); section.footer_distance = Cm(.65)
    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"; style.font.size = Pt(11)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.line_spacing = 1.07
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.widow_control = True
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if key.endswith("Theme"): del fonts.attrib[key]
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + key), "Times New Roman")
            for border in style.element.xpath("./w:pPr/w:pBdr"):
                border.getparent().remove(border)
    for name, size, before, after in (("Title", 16, 0, 6), ("Heading 1", 14, 0, 8), ("Heading 2", 11, 6, 5)):
        style = doc.styles[name]; style.font.size = Pt(size); style.font.bold = True
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "Bauwirtschaft · Werkstatt · Klotzkette"
    for run in header.runs: run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.text = "Redaktionsstand 28.09.2026 · Seite "
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    footer.add_run(" von 100")
    for run in footer.runs: run.font.size = Pt(9)
    cp = doc.core_properties; cp.author = "Klotzkette"; cp.last_modified_by = "Klotzkette"
    cp.title = TITLE; cp.subject = "Hundert verbundene Workflowstationen für Bauprojekte"
    cp.language = "de-DE"
    for number, title, content, _ in items:
        if number > 1: doc.add_page_break()
        if number == 1:
            doc.add_paragraph("Bauwirtschaft Werkstatt", "Title")
            p = doc.add_paragraph(INTRO)
            for run in p.runs: run.font.size = Pt(9)
        doc.add_paragraph(f"{number}. {title}", "Heading 1")
        for block in content.split("\n\n"):
            if block.startswith("### "):
                doc.add_paragraph(block[4:].strip(), "Heading 2")
            else:
                add_text(doc.add_paragraph(), block.replace("\n", " "))
    target = OUT / f"{STEM}.docx"; doc.save(target)
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--markdown-only", action="store_true")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--renderer", default=os.environ.get("DOCX_RENDERER"))
    parser.add_argument("--qa-dir", type=Path, default=Path("/tmp/bauwirtschaft-werkstatt-qa"))
    args = parser.parse_args(); items = sources(); md = markdown(items)
    if args.markdown_only:
        print(md.relative_to(ROOT)); return 0
    docx = word(items)
    result = {"source_sha256": source_digest(), "stations": 100, "markdown": str(md), "docx": str(docx)}
    if args.render:
        if not args.renderer or not Path(args.renderer).is_file():
            raise ValueError("Für native Sichtprüfung --renderer oder DOCX_RENDERER angeben")
        args.qa_dir.mkdir(parents=True, exist_ok=True)
        subprocess.run([sys.executable, args.renderer, str(docx), "--output_dir", str(args.qa_dir),
                        "--emit_pdf", "--width", "1000", "--height", "1500"], check=True)
        from pypdf import PdfReader
        rendered = args.qa_dir / f"{STEM}.pdf"; reader = PdfReader(rendered)
        if len(reader.pages) != 100:
            raise ValueError(f"Erwartet 100 redaktionell gesetzte Seiten, erhalten {len(reader.pages)}. Inhalt/Layout überarbeiten.")
        for number, page in enumerate(reader.pages, 1):
            if not re.search(rf"(?m)^{number}\.\s", page.extract_text() or ""):
                raise ValueError(f"Station {number} beginnt nicht auf der zugehörigen Seite")
        target = OUT / rendered.name; shutil.copyfile(rendered, target)
        result.update(pdf=str(target), pages=100, png_dir=str(args.qa_dir), visual_review="pending")
    (args.qa_dir).mkdir(parents=True, exist_ok=True)
    (args.qa_dir / "build.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2)); return 0


if __name__ == "__main__":
    raise SystemExit(main())
