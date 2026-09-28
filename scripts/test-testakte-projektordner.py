#!/usr/bin/env python3
"""Integration und Pfadsicherheit bestellter Projektordner-ZIPs. Autor: Klotzkette."""
from contextlib import redirect_stderr, redirect_stdout
import importlib.util
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from reportlab.pdfgen.canvas import Canvas
from testakte_file_filter import include_in_working_dump
from testakte_zip_common import safe_archive_name, structured_archive_pairs
from testakte_disclaimer import NOTICE_BYTES

SCRIPTS = Path(__file__).resolve().parent


def module(name, filename):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS/filename)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


W = module('projektordner_working', 'build-testakten-release-zips.py')
E = module('projektordner_pdf', 'build-testakten-einzelpdf-zips.py')
V = module('projektordner_validate', 'validate-testakten-release-zips.py')
P = module('projektordner_pdf_validate', 'validate-testakten-einzelpdf-zips.py')
G = module('projektordner_gesamt_validate', 'validate-testakten-gesamt-pdf.py')


class ProjectDirectories(unittest.TestCase):
    def test_only_named_diary_volume_is_allowed_and_required(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            cases = root/'testakten'
            case = cases/'bauwirtschaft-hildesheim-lebensakte'
            folder = case/'gesamt-pdf'
            folder.mkdir(parents=True)
            main = folder/(case.name+'_gesamt.pdf')
            diary = folder/(case.name+'_bautagebuch.pdf')
            for path in (main, diary):
                canvas = Canvas(str(path))
                canvas.drawString(50, 700, 'Baustellenbericht vom 1. Dezember 2027')
                canvas.save()
            (case/'README.md').write_text('\n'.join('gesamt-pdf/'+p.name for p in (main, diary)))
            with patch.object(G, 'ROOT', root), patch.object(G, 'TESTAKTEN', cases), redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                self.assertEqual(G.main(), 0)
                extra = folder/'nicht-bestellter-teilband.pdf'
                extra.write_bytes(main.read_bytes())
                self.assertEqual(G.main(), 1)
                extra.unlink()
                diary.unlink()
                self.assertEqual(G.main(), 1)

    def test_original_and_pdf_builders_keep_distinct_subfolders_and_exact_content(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            case = root/'bauwirtschaft-hildesheim-lebensakte'
            sources = {'08_Bauausfuehrung/Januar/bericht.txt': 'Am 7. Januar wurde die Schalung kontrolliert.',
                       '08_Bauausfuehrung/Februar/bericht.txt': 'Am 7. Februar wurde die Bewehrung kontrolliert.'}
            for relative, text in sources.items():
                path = case/relative
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(text, encoding='utf-8')
            dist = root/'dist'
            dist.mkdir()
            original, count = W.build_single(case, dist)
            self.assertEqual(count, 3)
            with zipfile.ZipFile(original) as archive:
                self.assertEqual(archive.namelist()[0], 'README.txt')
                self.assertEqual(set(archive.namelist()), {'README.txt', *sources})
                self.assertEqual(archive.read('README.txt'), NOTICE_BYTES)
                for path, text in sources.items():
                    self.assertEqual(archive.read(path), text.encode())
            self.assertEqual(len(V.zip_entries(original, require_notice=True, allow_directories=True)), 3)
            with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                V.zip_entries(original, require_notice=True)
            pdf, count = E.build_single(case, dist)
            self.assertEqual(count, 2)
            with zipfile.ZipFile(pdf) as archive:
                self.assertEqual(set(archive.namelist()), {'README.txt', *(str(Path(p).with_suffix('.pdf')) for p in sources)})
            self.assertEqual(len(P.zip_entries(pdf, expected_suffix='.pdf', allow_directories=True)), 3)

    def test_other_cases_remain_flat(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            case = root/'andere-akte'
            source = case/'eingang/brief.txt'
            source.parent.mkdir(parents=True)
            source.write_text('Bitte beachten Sie unseren Brief vom 7. Januar.', encoding='utf-8')
            dist = root/'dist'
            dist.mkdir()
            archive, _ = W.build_single(case, dist)
            with zipfile.ZipFile(archive) as result:
                self.assertEqual(result.namelist(), ['README.txt', 'eingang__brief.txt'])

    def test_path_traversal_absolute_windows_and_control_paths_are_rejected(self):
        for path in ('../brief.txt', '/brief.txt', '01/../../brief.txt', '01\\brief.txt', 'C:/brief.txt',
                     '01//brief.txt', './brief.txt', '01/brief.txt ', '01/CON.txt', '01/brief\n.txt',
                     '01/Datei*.txt', '01/Datei?.txt', '01/Datei<.txt', '01/Datei>.txt',
                     '01/Datei|.txt', '01/Datei".txt'):
            with self.subTest(path=path):
                self.assertFalse(safe_archive_name(path, allow_directories=True))
                with tempfile.TemporaryDirectory() as temporary:
                    archive = Path(temporary)/'unzulassig.zip'
                    with zipfile.ZipFile(archive, 'w') as output:
                        output.writestr('README.txt', NOTICE_BYTES)
                        output.writestr(path, b'Dokument')
                    with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                        V.zip_entries(archive, require_notice=True, allow_directories=True)

    def test_case_insensitive_collision_fails_instead_of_losing_an_original(self):
        with self.assertRaises(ValueError):
            structured_archive_pairs([(Path('a'), Path('Ordner/Brief.txt')), (Path('b'), Path('ordner/brief.txt'))])

    def test_requested_word_fields_are_allowed_but_answers_still_excluded(self):
        with tempfile.TemporaryDirectory() as temporary:
            case = Path(temporary)/'bauwirtschaft-hildesheim-lebensakte'
            folder = case/'12_Wordvorlagen/Bau'
            folder.mkdir(parents=True)
            form = folder/'Begehung.docx'
            document = Document()
            document.add_paragraph('Die Begehung findet am [Datum einsetzen] mit den benannten Beteiligten statt.')
            document.save(form)
            self.assertTrue(include_in_working_dump(form, case))
            wrong_place = case/'08_Bauausfuehrung'
            wrong_place.mkdir()
            actual = wrong_place/form.name
            actual.write_bytes(form.read_bytes())
            self.assertFalse(include_in_working_dump(actual, case))
            answer = folder/'Auswertung.docx'
            document.add_paragraph('Musterlösung und Erwartungshorizont')
            document.save(answer)
            self.assertFalse(include_in_working_dump(answer, case))


if __name__ == '__main__':
    unittest.main()
