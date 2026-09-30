#!/usr/bin/env python3
"""Native Unterlagen der beiden abgegrenzten SV-Arbeitsakten erzeugen."""
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
import json, csv, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from xml.sax.saxutils import escape

ROOT=Path(__file__).resolve().parents[1]
TMP=Path('/tmp/sozialversicherungspflicht-20260930/befreiung-akten')
TMP.mkdir(parents=True,exist_ok=True)
S='sozialversicherung-syndikus-versorgungswerk-hamburg'
A='sozialversicherung-ag-organe-hannover'
records=[]

def add(case,num,stem,kind,title,author,date,recipient,body):
    records.append(dict(case=case,num=num,stem=stem,kind=kind,title=title,author=author,date=date,recipient=recipient,body=body.strip()))

add(S,1,'Mila_an_Kanzlei','eml','Unterlagen zu Fleetbogen und meiner Altersversorgung','Mila Ahrens <mila.ahrens@postfach.example>','2026-09-28T18:42:00+02:00','Ottilie Semmel <kanzlei@semmel-weber.example>', '''Sehr geehrte Frau Semmel,

ich möchte wissen, ob die Abrechnung seit April richtig ist und was ich jetzt gegenüber dem Versorgungswerk, der Rentenversicherung und unserer Krankenkasse erklären muss. Unser Steuerbüro will wegen meines Wechsels zu Fleetbogen Netze sämtliche Monate ab April neu abrechnen. Ich hatte aber nie gekündigt und bekam bei dem Wechsel nicht einmal einen neuen Laptop. Seit dem Sommer mache ich allerdings deutlich mehr als vorher.

In der Anlage finden Sie den alten Arbeitsvertrag, die Übernahmevereinbarung und den Sommernachtrag. Die Kammer hatte im April geschrieben, dass meine Zulassung weiterbesteht. In der Personalakte steht deshalb noch der Rentenbescheid von 2019. Ich dachte, damit sei alles erledigt. Den elektronischen Antrag habe ich nach dem Gespräch mit Herrn Brehm am 18. September abgeschickt. Das Portalblatt liegt bei. Die Rentenversicherung möchte bis zum 8. Oktober weitere Unterlagen.

Bitte prüfen Sie auch, ob meine kleine eigene Kanzlei etwas ändert. Ich habe dafür kein Personal, arbeite am Wochenende zu Hause und habe die Mandate nicht von Fleetbogen erhalten. Mein Mann ist Lehrer; unsere Tochter wurde 2018 geboren. Ich bin freiwillig bei der Nordhafen Krankenkasse versichert. Ein Wechsel in die private Krankenversicherung war nie vereinbart.

Bitte senden Sie zunächst nur mir eine Einschätzung und Entwürfe. Mit dem Arbeitgeber möchte ich vor einer Meldung sprechen. Die Originale aus 2019 liegen im Schrank, die Worddateien entsprechen meinen unterschriebenen Ausfertigungen.

Mit freundlichen Grüßen
Mila Ahrens''')
add(S,2,'Personalangaben_Ahrens','docx','Angaben zur Person und zu meinen Tätigkeiten','Mila Ahrens','28. September 2026','Kanzlei Semmel und Weber, Hamburg', '''1 Angaben zur Person
Ich heiße Mila Ahrens, bin am 17. Mai 1988 geboren und wohne in der Falkenweide 18, 22305 Hamburg. Ich bin verheiratet und habe eine Tochter, geboren am 11. September 2018. Ich habe meinen Wohnsitz seit Beginn der Tätigkeit in Hamburg. Eine weitere Tätigkeit im Ausland übe ich nicht aus.

2 Beschäftigung
Seit dem 1. April 2019 arbeitete ich für die Fleetbogen Mobilität GmbH. Zum 1. April 2026 wurde mein Arbeitsverhältnis auf die Fleetbogen Netze GmbH übertragen. Mein Schreibtisch blieb in der Werftstraße 42. Seit dem 1. Juli 2026 lautet meine Funktionsbezeichnung Leiterin Recht und Geschäftsbetrieb. Im Arbeitsvertrag stehen 38 Wochenstunden. Im August waren es wegen der Einführung unseres Abrechnungssystems mehrfach mehr als 45 Stunden. Die Mehrstunden werden nicht getrennt ausgezahlt.

3 Versicherung und Versorgung
Ich bin Mitglied der Hanseatischen Rechtsanwaltskammer und des Versorgungswerks der Rechtsanwältinnen und Rechtsanwälte in Hamburg. Die Beiträge werden vom Gehalt einbehalten und zusammen mit dem Arbeitgeberanteil überwiesen. Meine Krankenkasse ist die Nordhafen Krankenkasse. Ihre Beitragsmitteilung nennt für 2026 einen Zusatzbeitrag von 3,2 Prozent. Eine private Pflegeversicherung habe ich nicht. Einen Bescheid zur Rentenversicherung für die seit April 2026 bestehende Arbeitgeberbezeichnung habe ich bislang nicht erhalten.

4 Weitere Einnahmen
Ich führe eine kleine Einzelkanzlei unter meinem Namen. Im Jahr 2025 betrugen deren Einnahmen 7.460 Euro, die Ausgaben 2.910 Euro. Für 2026 liegt noch kein Steuerbescheid vor. In der Kanzlei beschäftigen weder mein Mann noch ich Arbeitnehmer. Das Auftragsbuch habe ich beigefügt. Außerdem erhalten wir gemeinsam Miete für eine geerbte Garage. Mein Anteil betrug 2025 nach Kosten 840 Euro. Die Krankenkasse hat zuletzt im Februar den Steuerbescheid 2024 angefordert.

5 Unterlagen und Unterschrift
Meine Angaben zu den Arbeitsanteilen im Sommer sind Schätzungen. Eine Zeiterfassung nach Rechtsfragen und Betriebsführung gibt es nicht; der Kalender enthält jedoch einzelne Termine. Ich habe keine Erklärung unterzeichnet, nach der ich insgesamt auf Sozialversicherung verzichten wollte.

Hamburg, 28. September 2026
Mila Ahrens''')
add(S,3,'Arbeitsvertrag_2019','docx','Arbeitsvertrag für die Rechtsabteilung','Fleetbogen Mobilität GmbH und Mila Ahrens','18. März 2019','Ausfertigung für Mila Ahrens', '''Zwischen der Fleetbogen Mobilität GmbH, Werftstraße 42, 20457 Hamburg, vertreten durch Geschäftsführer Traugott Brehm, und Frau Mila Ahrens, Falkenweide 18, 22305 Hamburg, wird der folgende Arbeitsvertrag geschlossen.

1 Beginn und Aufgabe
Das Arbeitsverhältnis beginnt am 1. April 2019 und wird auf unbestimmte Zeit geschlossen. Frau Ahrens wird als Unternehmensjuristin in der Rechtsabteilung beschäftigt. Sie prüft Verträge über Ladeinfrastruktur, bearbeitet rechtliche Fragen des Vertriebs und vertritt die rechtlichen Interessen der Gesellschaft gegenüber Vertragspartnern. Die in der gesondert unterzeichneten Tätigkeitsbeschreibung vom heutigen Tag bezeichneten Aufgaben sind Bestandteil dieses Vertrags.

2 Fachliche Unabhängigkeit
Frau Ahrens erarbeitet ihre rechtliche Bewertung eigenverantwortlich. Fachliche Weisungen, die eine eigenständige Analyse der Rechtslage oder eine einzelfallbezogene Rechtsberatung ausschließen, werden nicht erteilt. Kaufmännische Entscheidungen trifft die Geschäftsführung nach Erhalt der rechtlichen Empfehlung. Frau Ahrens darf auf rechtliche Bedenken schriftlich hinweisen und diese unverändert zur Entscheidung vorlegen. Sie erhält keine Befugnis, eigenständig Produkte, Preise oder Personalbudgets festzulegen.

3 Arbeitszeit und Arbeitsort
Die regelmäßige Arbeitszeit beträgt 38 Stunden je Woche, verteilt auf Montag bis Freitag. Beginn und Ende werden innerhalb der betrieblichen Gleitzeit abgestimmt. Arbeitsort ist die Werftstraße 42 in Hamburg. Bis zu zwei Tage pro Woche können nach Abstimmung mobil gearbeitet werden. Bei dringenden Verhandlungen kann die Gesellschaft angemessene Mehrarbeit anordnen; sie wird innerhalb von drei Monaten durch Freizeit ausgeglichen. Gesetzliche Arbeitszeitgrenzen bleiben maßgeblich.

4 Vergütung und Auslagen
Die monatliche Bruttovergütung beträgt 5.800 Euro und wird spätestens am letzten Bankarbeitstag des Monats gezahlt. Ein Anspruch auf erfolgsabhängige Vergütung besteht nicht. Notwendige Dienstreisekosten werden gegen Beleg nach der jeweils mitgeteilten Reisekostenordnung erstattet. Die Gesellschaft führt Steuern und gesetzlich geschuldete Beiträge ab. Ein Zuschuss zur berufsständischen Versorgung wird nach Vorlage der hierfür erforderlichen Unterlagen geleistet.

5 Urlaub und Verhinderung
Frau Ahrens erhält bei einer Fünftagewoche 30 Arbeitstage Urlaub im Kalenderjahr. Die zeitliche Lage ist mit der Leitung der Rechtsabteilung abzustimmen. Bei Arbeitsunfähigkeit informiert sie die Gesellschaft unverzüglich über die voraussichtliche Dauer. Die Vergütung wird nach den gesetzlichen Bestimmungen fortgezahlt. Persönliche Daten und Krankheitsdiagnosen sind der Geschäftsführung nicht unaufgefordert mitzuteilen.

6 Berufliche Zulassung und Nebentätigkeit
Die Gesellschaft unterstützt den Antrag auf Zulassung als Syndikusrechtsanwältin für diese Beschäftigung und bestätigt die fachliche Unabhängigkeit. Frau Ahrens zeigt ihr den Eingang einer Zulassungs- oder Befreiungsentscheidung an. Ihre selbständige anwaltliche Nebentätigkeit für eigene Mandanten wird gestattet, soweit keine Interessenkollision, kein Wettbewerb mit der Gesellschaft und keine Beeinträchtigung der Arbeitsleistung entsteht. Mandantendaten werden nicht in den Systemen der Gesellschaft gespeichert.

7 Verschwiegenheit und Arbeitsergebnisse
Geschäftliche und personenbezogene Informationen dürfen nur zur Erfüllung der jeweiligen Aufgaben verwendet werden. Die Verschwiegenheitspflicht besteht nach Vertragsende fort. Die Gesellschaft erhält an im Rahmen der Beschäftigung erstellten Vertragsmustern und internen Schulungsunterlagen die für ihren Geschäftsbetrieb erforderlichen Nutzungsrechte. Berufsrechtliche Pflichten der Rechtsanwältin bleiben unberührt.

8 Kündigung und Änderungen
Nach einer Probezeit von sechs Monaten kann das Arbeitsverhältnis mit einer Frist von drei Monaten zum Monatsende gekündigt werden. Verlängern zwingende Vorschriften die Arbeitgeberfrist, gilt diese Verlängerung auch für die Arbeitnehmerin. Das Recht zur außerordentlichen Kündigung bleibt unberührt. Kündigungen bedürfen der gesetzlichen Schriftform. Änderungen sollen schriftlich dokumentiert werden; individuelle Abreden bleiben wirksam. Es bestehen keine Nebenabreden.

Hamburg, 18. März 2019
Traugott Brehm, Geschäftsführer
Mila Ahrens''')
add(S,4,'Taetigkeitsbeschreibung_2019','docx','Tätigkeitsbeschreibung der Unternehmensjuristin','Fleetbogen Mobilität GmbH','18. März 2019','Hanseatische Rechtsanwaltskammer', '''1 Gegenstand der Beschäftigung
Frau Mila Ahrens soll Rechtsfragen zu Einkauf, Bau und Betrieb gewerblicher Ladepunkte bearbeiten. Sie klärt die tatsächlichen Abläufe im Gespräch mit Projektleitern auf, prüft das Vertrags- und Haftungsrisiko und entwickelt rechtliche Handlungsalternativen. Ihr unmittelbarer Vorgesetzter ist Geschäftsführer Traugott Brehm. Eine fachliche Vorprüfung ihrer rechtlichen Bewertung durch den Vertrieb findet nicht statt.

2 Beratung und Gestaltung
Frau Ahrens entwirft Rahmenverträge, prüft Änderungswünsche und berät die Fachabteilungen über die rechtlichen Folgen. Sie führt Verhandlungen über Gewährleistung, Kündigung, Sicherheiten und Datenschutz. Sie darf Erklärungen im Namen der Gesellschaft abgeben, soweit die Geschäftsführung die jeweilige wirtschaftliche Entscheidung freigegeben hat. Die Freigabe betrifft den Geschäftsabschluss, nicht den Inhalt ihrer rechtlichen Bewertung.

3 Arbeitsumfang
Nach der bei Einstellung vorgesehenen Verteilung entfallen etwa 45 Prozent der Arbeitszeit auf Vertragsgestaltung, 25 Prozent auf Beratung und Verhandlungen, 20 Prozent auf Streitbearbeitung und zehn Prozent auf interne Schulungen und Dokumentation. Diese Verteilung beschreibt die geplante Stelle und ist keine minutengenaue Abrechnung. Frau Ahrens übernimmt weder die Einsatzplanung der Monteure noch die Verantwortung für Mahnläufe oder laufende Kundendienstabrechnungen.

4 Unabhängigkeit und Vertretung
Die Gesellschaft verpflichtet sich, Frau Ahrens die fachlich unabhängige und eigenverantwortliche Ausübung dieser Aufgaben zu ermöglichen. Wirtschaftliche Zielvorgaben dürfen ihre eigenständige Prüfung und Darstellung der Rechtslage nicht ersetzen. Während ihrer Abwesenheit übernimmt Justitiarin Elfriede Runge die Rechtsfragen. Außerhalb des hier beschriebenen Arbeitsverhältnisses erfolgt keine Beratung von Kunden als deren Rechtsanwältin.

Hamburg, 18. März 2019
Traugott Brehm
Mila Ahrens''')
add(S,5,'Zulassung_Syndikus_2019','pdf','Zulassung zur Rechtsanwaltschaft als Syndikusrechtsanwältin','Hanseatische Rechtsanwaltskammer','28. Mai 2019','Frau Mila Ahrens, Falkenweide 18, 22305 Hamburg', '''Aktenzeichen ZS 219 19

Sehr geehrte Frau Ahrens,

1 Entscheidung
Sie werden für Ihre Tätigkeit bei der Fleetbogen Mobilität GmbH auf der Grundlage des Arbeitsvertrags und der Tätigkeitsbeschreibung vom 18. März 2019 als Syndikusrechtsanwältin zugelassen. Ihr Antrag ist am 25. März 2019 bei uns eingegangen. Nach der Arbeitgeberbestätigung haben Sie die Tätigkeit am 1. April 2019 aufgenommen. Ihre Kammermitgliedschaft für diese Zulassung beginnt mit diesem Tag.

2 Umfang der Entscheidung
Gegenstand der Prüfung sind die anwaltlichen Tätigkeiten in der Rechtsabteilung des im Antrag bezeichneten Arbeitgebers. Die Deutsche Rentenversicherung Bund wurde angehört. Der Arbeitsvertrag gewährleistet nach den vorgelegten Unterlagen die fachliche Unabhängigkeit. Weitere Arbeitsverhältnisse sind nicht Gegenstand dieser Entscheidung. Ihre bereits bestehende Zulassung als niedergelassene Rechtsanwältin wird hiervon nicht berührt.

3 Hinweise und Rechtsschutz
Bitte zeigen Sie tätigkeitsbezogene Änderungen des Arbeitsvertrags und wesentliche Änderungen der Tätigkeit unverzüglich an und reichen Sie die entsprechenden Unterlagen ein. Die sozialversicherungsrechtliche Befreiungsentscheidung wird von dieser Kammer nicht getroffen. Gegen diesen Bescheid kann innerhalb eines Monats nach Zustellung Klage beim Anwaltsgerichtshof erhoben werden. Maßgeblich sind die gesetzlichen Anforderungen an Form und Übermittlung der Klage.

Mit freundlichen Grüßen
Elisabeth Fenn, Zulassungsabteilung
Zugangsvermerk der Empfängerin: Brief erhalten am 31. Mai 2019.''')
add(S,6,'Versorgungswerk_Aufnahme','pdf','Mitgliedschaft und Beitragskonto','Versorgungswerk der Rechtsanwältinnen und Rechtsanwälte in Hamburg','6. Juni 2019','Mila Ahrens', '''Mitgliedsnummer 48192

Sehr geehrte Frau Ahrens,

wir bestätigen Ihre Pflichtmitgliedschaft seit dem 1. April 2019 auf der Grundlage der von der Rechtsanwaltskammer übermittelten Daten. Ihr Beitragskonto wird unter der oben genannten Mitgliedsnummer geführt. Bitte verwenden Sie diese Nummer bei Zahlungen und bei Nachrichten an die Mitgliederverwaltung.

Für Ihre Beschäftigung bei der Fleetbogen Mobilität GmbH haben wir die schriftlich eingereichte Arbeitgeberinformation in Ihrer Mitgliederakte vermerkt. Die Festsetzung der laufenden Beiträge erhalten Sie gesondert. Eine Änderung des Arbeitseinkommens, ein Arbeitgeberwechsel und Änderungen der beruflichen Tätigkeit sind uns mitzuteilen. Eine laufende Beitragszahlung ersetzt eine erforderliche Anzeige nicht.

Ihr Antrag auf Befreiung von der gesetzlichen Rentenversicherungspflicht ist bei uns am 12. April 2019 eingegangen. Wir haben die Mitgliedschaftsangaben bestätigt und den Antrag zur Entscheidung weitergeleitet. Über die Befreiung entscheidet die Deutsche Rentenversicherung Bund. Bitte legen Sie uns eine Ausfertigung der Entscheidung vor, sobald sie Ihnen vorliegt.

Für Rückfragen zu Ihrem Beitragskonto schreiben Sie bitte an mitglieder@versorgung-hamburg.example. Geben Sie nur die für die Bearbeitung erforderlichen Unterlagen an.

Mit freundlichen Grüßen
Hildegard Ulmer, Mitgliederverwaltung''')
add(S,7,'DRV_Befreiungsbescheid_2019','pdf','Befreiung von der Rentenversicherungspflicht','Deutsche Rentenversicherung Bund','22. Juli 2019','Frau Mila Ahrens', '''Versicherungsnummer 12 170588 A 507
Geschäftszeichen BF 2019 48192

Sehr geehrte Frau Ahrens,

1 Entscheidung
Auf Ihren am 12. April 2019 eingegangenen Antrag werden Sie nach Paragraf 6 Absatz 1 Satz 1 Nummer 1 SGB VI ab dem 1. April 2019 von der Versicherungspflicht in der gesetzlichen Rentenversicherung befreit.

Die Entscheidung betrifft Ihre Beschäftigung als Syndikusrechtsanwältin in der Rechtsabteilung der Fleetbogen Mobilität GmbH, Werftstraße 42, 20457 Hamburg. Grundlage sind der Arbeitsvertrag und die Tätigkeitsbeschreibung vom 18. März 2019 sowie die Zulassungsentscheidung der Hanseatischen Rechtsanwaltskammer vom 28. Mai 2019.

2 Gründe und Mitteilungspflichten
Die Pflichtmitgliedschaften in der berufsständischen Kammer und Versorgungseinrichtung wurden bestätigt. Ihr Antrag wurde innerhalb von drei Monaten nach Vorliegen der Befreiungsvoraussetzungen gestellt. Die Befreiung ist auf die oben bezeichnete Beschäftigung beschränkt. Änderungen dieser Beschäftigung oder ihrer tatsächlichen Ausgestaltung teilen Sie uns bitte mit. Die Entscheidung trifft keine Feststellung zur Kranken-, Pflege-, Arbeitslosen- oder Unfallversicherung.

3 Rechtsbehelf
Gegen diesen Bescheid können Sie innerhalb eines Monats nach Bekanntgabe Widerspruch erheben. Der Widerspruch ist bei der Deutschen Rentenversicherung Bund schriftlich, in zugelassener elektronischer Form oder zur Niederschrift einzureichen. Bitte geben Sie dabei das Geschäftszeichen an.

Dieser Bescheid wurde maschinell erstellt und ist ohne Unterschrift gültig.
Vermerk Personalstelle vom 29. Juli 2019: Kopie zur Entgeltakte genommen.''')
add(S,8,'Dreiseitige_Uebernahme','docx','Vereinbarung über die Übernahme des Arbeitsverhältnisses','Fleetbogen Mobilität GmbH, Fleetbogen Netze GmbH und Mila Ahrens','16. März 2026','Je eine Ausfertigung für die drei Vertragsparteien', '''Die Fleetbogen Mobilität GmbH, Werftstraße 42, 20457 Hamburg, vertreten durch Traugott Brehm, die Fleetbogen Netze GmbH, ebenfalls Werftstraße 42, vertreten durch Juna Martens, und Frau Mila Ahrens vereinbaren die Übernahme ihres bestehenden Arbeitsverhältnisses.

1 Übernahme
Mit Wirkung zum 1. April 2026 tritt die Fleetbogen Netze GmbH anstelle der Fleetbogen Mobilität GmbH in das seit dem 1. April 2019 bestehende Arbeitsverhältnis mit Frau Ahrens ein. Das Arbeitsverhältnis wird mit sämtlichen Rechten und Pflichten unverändert fortgeführt. Es wird weder gekündigt noch neu begründet. Frau Ahrens stimmt dem Wechsel der Arbeitgeberin zu. Diese Vereinbarung enthält keine Festlegung, ob unabhängig davon die gesetzlichen Voraussetzungen eines Betriebsübergangs erfüllt sind.

2 Vertragsgrundlagen und Betriebszugehörigkeit
Der Arbeitsvertrag vom 18. März 2019, die damalige Tätigkeitsbeschreibung und der Vergütungsnachtrag vom 12. Dezember 2024 gelten fort. Als Beginn der Betriebszugehörigkeit bleibt der 1. April 2019 maßgeblich. Eine neue Probezeit wird nicht vereinbart. Alle zum Übernahmestichtag bestehenden Urlaubsansprüche und Zeitguthaben gehen auf die übernehmende Gesellschaft über. Der Urlaubskontostand wird nach Abschluss der Märzabrechnung mitgeteilt.

3 Tätigkeit und Arbeitsbedingungen
Frau Ahrens arbeitet weiterhin als Syndikusrechtsanwältin in der Rechtsabteilung. Ihre fachliche Unabhängigkeit bleibt entsprechend Ziffer 2 des Ausgangsvertrags gewährleistet. Arbeitsort, regelmäßige Arbeitszeit, mobile Arbeit und Vertretung ändern sich durch die Übernahme nicht. Die monatliche Bruttovergütung beträgt aufgrund des fortgeltenden Nachtrags 6.500 Euro. Eine spätere Änderung der Funktion bedarf einer gesonderten Vereinbarung.

4 Offene Ansprüche
Die übernehmende Gesellschaft erfüllt auch die vor dem Stichtag entstandenen, noch nicht erfüllten Ansprüche aus dem Arbeitsverhältnis. Die bisherige Gesellschaft wird im Innenverhältnis zum Stichtag abgerechnet. Gesetzliche Haftungsansprüche von Frau Ahrens werden durch diese interne Abrechnung nicht beschränkt. Die Beteiligten verzichten nicht auf unbekannte Vergütungs-, Urlaubs- oder Erstattungsansprüche.

5 Unterlagen und Personalverwaltung
Die für die Fortführung erforderliche Personalakte wird unter Beachtung der Datenschutzvorschriften an die neue Arbeitgeberin übergeben. Frau Ahrens erhält eine Kopie des Übernahmeprotokolls. Die neue Arbeitgeberin übernimmt die laufenden Abrechnungsdaten. Berufsrechtlich oder versicherungsrechtlich erforderliche Anzeigen werden von den jeweils zuständigen Beteiligten vorgenommen; diese Vereinbarung ersetzt keine behördliche Entscheidung.

6 Sonstige Abreden
Die genehmigte eigene anwaltliche Nebentätigkeit bleibt gestattet. Eine Ausgleichsquittung, ein Verzicht auf Ansprüche oder eine neue Wettbewerbsklausel ist mit der Übernahme nicht verbunden. Änderungen dieser Vereinbarung werden zu Beweiszwecken schriftlich dokumentiert. Individuelle Abreden und gesetzliche Formvorschriften bleiben unberührt.

Hamburg, 16. März 2026
Traugott Brehm für die Fleetbogen Mobilität GmbH
Juna Martens für die Fleetbogen Netze GmbH
Mila Ahrens''')
add(S,9,'HR_Uebernahme_Rundmail','eml','Personalnummer bleibt gleich','Traugott Brehm <t.brehm@fleetbogen.example>','2026-03-24T10:16:00+01:00','Mila Ahrens <m.ahrens@fleetbogen.example>', '''Liebe Mila,

die unterschriebene Übernahme ist vollständig zurück. Deine Personalnummer 0047 bleibt auch bei Netze bestehen. Unser Dienstleister legt technisch einen neuen Mandanten an, übernimmt aber Eintrittsdatum, Urlaub und Zeitkonto. Bitte wundere Dich deshalb nicht über eine Abmeldung und Anmeldung in den nächsten Wochen. Das ist die technische Verarbeitung des Gesellschaftswechsels.

Ich habe dem Dienstleister Deinen alten Rentenbescheid geschickt. Er fragte, ob es einen neuen gibt. Nach meinem Verständnis hat sich an Deinem Job nichts geändert; ich habe ihm deshalb zunächst geschrieben, dass er die bisherige Versorgung fortführen soll. Das war keine Abstimmung mit der Rentenversicherung. Bitte schicke die Vereinbarung auch an die Kammer, damit unsere Akte vollständig ist.

Im Sommer müssen wir über die Leitung des Geschäftsbetriebs reden. Juna kann den Kundendienst nicht dauerhaft nebenbei führen. Die Übernahmevereinbarung enthält dazu noch nichts, und das vorläufige Organigramm ist nicht als Vertragsänderung gedacht.

Viele Grüße
Traugott''')
add(S,10,'Anzeige_RAK_April','eml','Übernahme meines Arbeitsverhältnisses zum 1. April','Mila Ahrens <mila.ahrens@postfach.example>','2026-04-02T20:11:00+02:00','Zulassungen <zulassung@rak-hamburg.example>', '''Sehr geehrte Damen und Herren,

ich übersende die von allen drei Parteien unterzeichnete Vereinbarung vom 16. März 2026. Die Fleetbogen Netze GmbH hat mein Arbeitsverhältnis zum 1. April mit sämtlichen Rechten und Pflichten übernommen. Ich bitte um Mitteilung, ob meine Zulassung für die fortgeführte Tätigkeit entsprechend dokumentiert werden kann.

Nach meiner derzeitigen Tätigkeit und der Vereinbarung bleiben Aufgaben, Arbeitszeit und fachliche Unabhängigkeit unverändert. Ich bearbeite die bisherigen Verträge und Streitfälle. Die kaufmännische Verantwortung liegt weiterhin bei Frau Martens. Für den Sommer wird intern über eine andere Aufgabenverteilung gesprochen; es gibt hierzu noch keine von mir unterzeichnete Änderung. Sollte eine solche Änderung vereinbart werden, reiche ich die Unterlagen nach.

Beigefügt sind die Vereinbarung und die Arbeitgeberbestätigung über die unveränderte Tätigkeit als Worddateien. Bitte bestätigen Sie den Eingang. Meine Mitgliedsnummer lautet H 80461. Meine private Kanzlei wird daneben unverändert fortgeführt.

Mit freundlichen Grüßen
Mila Ahrens''')
add(S,11,'RAK_Fortgeltung_April','pdf','Feststellung zur fortgeführten Beschäftigung','Hanseatische Rechtsanwaltskammer','27. April 2026','Frau Rechtsanwältin Mila Ahrens', '''Aktenzeichen ZS 219 19 F

Sehr geehrte Frau Ahrens,

auf Ihre Anzeige vom 2. April 2026 stellen wir fest, dass Ihre Zulassung als Syndikusrechtsanwältin die bei der Fleetbogen Netze GmbH fortgeführte Tätigkeit umfasst. Die Übernahmevereinbarung vom 16. März 2026 überträgt das Arbeitsverhältnis mit sämtlichen Rechten und Pflichten. Nach Ihrer Erklärung und der Arbeitgeberbestätigung ändern sich die für die Zulassung maßgeblichen Tätigkeiten nicht. Die Deutsche Rentenversicherung Bund erhielt Gelegenheit zur Stellungnahme.

Diese Feststellung beruht auf der zum 1. April 2026 beschriebenen Fortführung. Die in Ihrer Nachricht erwähnte, noch nicht vereinbarte Neuordnung des Geschäftsbetriebs ist nicht Gegenstand dieses Bescheids. Bitte zeigen Sie eine entsprechende Vertrags- oder Tätigkeitsänderung unverzüglich unter Vorlage der konkreten Unterlagen an.

Die im Jahr 2019 festgestellte fachliche Unabhängigkeit muss weiterhin vertraglich und tatsächlich gewährleistet sein. Dieser Bescheid ersetzt keine Entscheidung eines Sozialversicherungsträgers über eine künftig geänderte Tätigkeit.

Gegen diese Entscheidung kann innerhalb eines Monats nach Zustellung Klage beim zuständigen Anwaltsgerichtshof erhoben werden. Die gesetzlichen Anforderungen an Form und Übermittlung gelten auch bei elektronischer Einreichung.

Mit freundlichen Grüßen
Elisabeth Fenn
Eingang bei der Empfängerin: 30. April 2026''')
add(S,12,'Sommer_Arbeitsplanung','docx','Arbeitsplanung für Recht und Geschäftsbetrieb','Juna Martens, Geschäftsführerin Fleetbogen Netze GmbH','10. Juni 2026','Mila Ahrens und Elfriede Runge', '''1 Ausgangslage
Im Kundendienst sind seit März zwei Leitungsstellen unbesetzt. Die Ticketzahl steigt, weil die neue Ladekartenplattform im Mai eingeführt wurde. Frau Ahrens kennt die Verträge und soll ab Juli die Abstimmung zwischen Kundendienst, Abrechnung und Rechtsabteilung übernehmen. Frau Runge bleibt für die laufenden gerichtlichen Verfahren zuständig.

2 Geplante Aufgaben
Frau Ahrens soll die Monatsziele für die Bearbeitung offener Tickets festlegen, die Teamleiter wöchentlich sprechen und die Freigabe größerer Erstattungen koordinieren. Bei strittigen Vertragsfragen soll sie weiterhin selbst rechtlich beraten. Für rein technische Störungen bleibt Herr Friedrich Knoop verantwortlich. Die budgetrelevanten Entscheidungen werden bis zu einer schriftlichen Neuregelung der Geschäftsführerin vorgelegt.

3 Zeitlicher Ablauf
Am 15. Juni beginnt eine zweiwöchige Übergabe mit Herrn Knoop. In dieser Phase soll Frau Ahrens an den Morgenrunden teilnehmen und die Vorgänge kennenlernen. Die formale Funktionsübernahme ist für den 1. Juli vorgesehen. Frau Ahrens hat in der Besprechung darauf hingewiesen, dass sie vor der endgültigen Aufgabenverteilung ihre berufsrechtlichen Unterlagen prüfen möchte. Ein Gespräch mit der Kammer ist noch nicht geführt worden.

4 Personal und Vertretung
Die sechs Mitarbeiter des Kundendiensts berichten fachlich an die jeweiligen Teamleiter und organisatorisch an Frau Ahrens. Die Geschäftsführerin entscheidet über Einstellungen und Kündigungen. Frau Ahrens soll Urlaubspläne koordinieren und Vorschläge zu Schulungen machen. Für die juristische Arbeit wird eine studentische Hilfskraft gesucht. Bis zu deren Einstellung muss die Vertragsablage gemeinsam betreut werden.

Juna Martens
Verteilt im Leitungsmeeting am 10. Juni 2026''')
add(S,13,'Nachtrag_Leitung_Juli','docx','Nachtrag zum fortgeführten Arbeitsvertrag','Fleetbogen Netze GmbH und Mila Ahrens','24. Juni 2026','Ausfertigung für beide Vertragsparteien', '''Die Fleetbogen Netze GmbH, vertreten durch Geschäftsführerin Juna Martens, und Frau Mila Ahrens ändern den mit Vereinbarung vom 16. März 2026 übernommenen Arbeitsvertrag vom 18. März 2019 wie folgt.

1 Funktion und Beginn
Frau Ahrens übernimmt mit Wirkung zum 1. Juli 2026 die Funktion Leiterin Recht und Geschäftsbetrieb. Sie bleibt für die juristische Beratung der Gesellschaft zuständig und übernimmt zusätzlich die organisatorische Leitung von Kundendienst und Abrechnung. Die neue Funktion begründet keine Bestellung zur Geschäftsführerin und keine Prokura.

2 Rechtsberatung
Die fachliche Unabhängigkeit bei der Prüfung und Beratung in Rechtsfragen gemäß Ziffer 2 des Ausgangsvertrags bleibt bestehen. Frau Ahrens entscheidet eigenverantwortlich, wie sie die Rechtslage bewertet und welche rechtlichen Handlungsmöglichkeiten sie empfiehlt. Wirtschaftliche Vorgaben verpflichten sie nicht zu einer hiervon abweichenden rechtlichen Stellungnahme. Ihre rechtlichen Bedenken werden unverändert an die Geschäftsführung weitergegeben.

3 Geschäftsbetrieb
Frau Ahrens legt im Rahmen des von der Geschäftsführung freigegebenen Jahresbudgets die operative Bearbeitungsreihenfolge fest. Sie darf Kunden bis zu 15.000 Euro je Einzelfall Erstattungen aus kaufmännischen Gründen bewilligen. Bei diesen Entscheidungen beachtet sie die wöchentlichen Vorgaben der Geschäftsführerin zu Kosten und Servicezielen. Sie ist für die Umsetzung der vorgegebenen Kennzahlen verantwortlich und berichtet jeweils montags. Einstellungen, Kündigungen und die Änderung des Gesamtbudgets bleiben der Geschäftsführung vorbehalten.

4 Vergütung und Arbeitszeit
Die monatliche Bruttovergütung steigt ab dem 1. Juli 2026 auf 8.000 Euro. Die regelmäßige Wochenarbeitszeit beträgt weiterhin 38 Stunden. Die Regelungen zum Freizeitausgleich gelten fort. Eine Zielprämie wird nicht vereinbart. Dienstreisekosten werden weiterhin gegen Beleg erstattet.

5 Berufliche Unterlagen
Frau Ahrens wird der Rechtsanwaltskammer die geänderte Aufgabenbeschreibung vorlegen. Die Gesellschaft stellt hierfür die tatsächlichen organisatorischen Abläufe und Befugnisse dar. Die Parteien geben mit diesem Nachtrag keine gemeinsame Erklärung über die beitragsrechtliche Behandlung ab. Soweit weitere Nachweise von einer zuständigen Stelle angefordert werden, unterstützen sich die Parteien bei deren Beschaffung.

6 Fortgeltung
Im Übrigen gelten der Ausgangsvertrag und die Übernahmevereinbarung unverändert fort. Insbesondere bleiben Urlaub, Kündigungsfristen und die genehmigte anwaltliche Nebentätigkeit bestehen. Die Wirkung dieses Nachtrags wird nicht davon abhängig gemacht, wann interne Stellenbezeichnungen im Personalprogramm aktualisiert werden.

Hamburg, 24. Juni 2026
Juna Martens
Mila Ahrens''')
add(S,14,'Personalstelle_alter_Bescheid','eml','Rückfrage zum alten Befreiungsbescheid','Traugott Brehm <t.brehm@fleetbogen.example>','2026-09-04T09:27:00+02:00','Mila Ahrens <m.ahrens@fleetbogen.example>', '''Liebe Mila,

unser Abrechner hat bei der Durchsicht den Aprilwechsel und Deinen Julinachtrag nebeneinandergelegt. Er möchte einen Bescheid, in dem Fleetbogen Netze ausdrücklich genannt ist. Ich habe nur die Kammerentscheidung vom April und den Rentenbescheid von 2019. Bisher steht im System weiterhin die berufsständische Versorgung.

Ich hatte im Frühjahr angenommen, dass ein konzerninterner Wechsel ohne neue Probezeit nichts verändert. Nachlesen kann ich in unseren Mails aber nur meine eigene Nachricht an den Dienstleister. Bitte verlasse Dich also nicht darauf, dass wir damals eine Zusage der Rentenversicherung eingeholt hätten. Eine solche Zusage finde ich nicht.

Die Meldungen möchte der Dienstleister noch nicht korrigieren, bis Du die Unterlagen mit Deiner Beraterin besprochen hast. Er hat den 8. Oktober als internen Wiedervorlagetermin gesetzt. Bitte reiche den Julinachtrag bei der Kammer ein und teile uns mit, was für die Septemberabrechnung benötigt wird. Den Nachweis über Deinen Antrag können wir zur Akte nehmen; ob er bereits einen Bescheid ersetzt, weiß ich nicht.

Viele Grüße
Traugott''')
add(S,15,'Antrag_Erstreckung_RAK','docx','Antrag zur geänderten Tätigkeit bei Fleetbogen Netze','Rechtsanwältin Mila Ahrens','16. September 2026','Hanseatische Rechtsanwaltskammer, Zulassungsabteilung', '''Sehr geehrte Damen und Herren,

ich beantrage, meine Zulassung als Syndikusrechtsanwältin auf die seit dem 1. Juli 2026 geänderte Tätigkeit bei der Fleetbogen Netze GmbH zu erstrecken. Die Gesellschaft hat mir zusätzlich die Leitung von Kundendienst und Abrechnung übertragen. Ich lege den Nachtrag vom 24. Juni und die Arbeitsplanung vom 10. Juni bei.

1 Tatsächliche Ausübung
Seit Anfang Juli nehme ich an den wöchentlichen Kennzahlenbesprechungen teil, entscheide über Erstattungen bis 15.000 Euro und koordiniere die sechs Mitarbeiter über deren Teamleiter. Mein Kalender enthält außerdem Vertragsverhandlungen, rechtliche Stellungnahmen und die Abstimmung mit externen Prozessbevollmächtigten. Eine nach Aufgaben getrennte Zeiterfassung führe ich nicht. Für die letzten vier Wochen schätze ich den juristischen Anteil auf etwas mehr als die Hälfte. Im Juli nahm die Einführung der Abrechnungssoftware erheblich mehr Zeit in Anspruch.

2 Fachliche Unabhängigkeit
Die Vertragsregelung zur eigenständigen Rechtsberatung wurde beibehalten. Frau Martens entscheidet über das Budget und gibt mir operative Ziele vor. Sie hat keine Änderung einer rechtlichen Stellungnahme verlangt. Wenn ich über einen Kulanzbetrag entscheide, berücksichtige ich allerdings auch Bindungsdauer, Kundenwert und die jeweils vorgegebenen Kostenlimits.

3 Frühere Mitteilung
Ihre Entscheidung vom 27. April betraf die unveränderte Fortsetzung nach der dreiseitigen Vertragsübernahme. Ich hatte angekündigt, eine spätere Änderung nachzureichen. Das ist im Juli wegen des Projektstarts unterblieben. Die Arbeitgeberbestätigung zur nunmehr ausgeübten Tätigkeit wird nachgereicht, sobald Frau Martens sie unterzeichnet hat. Bitte teilen Sie mir mit, ob weitere konkrete Angaben erforderlich sind.

Mit freundlichen Grüßen
Mila Ahrens''')
add(S,16,'Befreiungsantrag_Eingang','pdf','Eingangsbestätigung zum elektronischen Befreiungsantrag','Versorgungswerk der Rechtsanwältinnen und Rechtsanwälte in Hamburg','18. September 2026','Mila Ahrens, Mitgliedsnummer 48192', '''Übermittlung am 18. September 2026 um 21:14:38 Uhr
Vorgangsnummer EB 2026 48192 0918

Ihr elektronischer Antrag auf Befreiung von der gesetzlichen Rentenversicherungspflicht wurde entgegengenommen. Die nachfolgenden Angaben wurden übermittelt.

Arbeitgeber: Fleetbogen Netze GmbH, Werftstraße 42, 20457 Hamburg.
Beginn des Arbeitsverhältnisses bei diesem Arbeitgeber: 1. April 2026.
Bezeichnung der Tätigkeit: Leiterin Recht und Geschäftsbetrieb, Syndikusrechtsanwältin.
Beantragter Beginn der Befreiung: 1. April 2026.
Bisheriger Befreiungsbescheid: Deutsche Rentenversicherung Bund vom 22. Juli 2019, Geschäftszeichen BF 2019 48192.

Als Anlagen wurden die Übernahmevereinbarung, die Kammerentscheidung vom 27. April 2026 und der Vertragsnachtrag vom 24. Juni 2026 hochgeladen. Die Antragstellerin hat angegeben, dass ein Antrag zur geänderten Tätigkeit bei der Kammer gestellt wurde und die Entscheidung hierzu noch aussteht.

Diese Bestätigung dokumentiert den Eingang. Sie enthält keine Entscheidung über die gesetzlichen Voraussetzungen, den Umfang oder den Beginn einer Befreiung. Nach Abschluss der erforderlichen Bestätigungen wird der Vorgang zur Entscheidung an die Deutsche Rentenversicherung Bund weitergeleitet. Bewahren Sie die Vorgangsnummer für Rückfragen auf.

Automatisch erzeugte Bestätigung des Mitgliederportals''')
add(S,17,'DRV_Unterlagenanforderung','eml','Ihr Antrag EB 2026 48192 0918','Sachbearbeitung Befreiung <befreiung@drv-post.example>','2026-09-25T11:03:00+02:00','Mila Ahrens <mila.ahrens@postfach.example>', '''Sehr geehrte Frau Ahrens,

zu Ihrem über die Versorgungseinrichtung übermittelten Antrag benötigen wir noch Angaben über den Inhalt und die zeitliche Entwicklung Ihrer Beschäftigung. Bitte reichen Sie bis zum 8. Oktober 2026 die aktuelle Arbeitgeberbestätigung, die Entscheidung der Rechtsanwaltskammer zur Tätigkeit ab Juli und eine Beschreibung der tatsächlich ausgeübten Aufgaben ein. Sollte die Kammer noch nicht entschieden haben, genügt zunächst eine Mitteilung über den Bearbeitungsstand.

Bitte erläutern Sie insbesondere, ob das Arbeitsverhältnis im April vollständig übernommen wurde oder zuvor beendet worden war. Die von Ihnen vorgelegte Vereinbarung nennt eine Fortführung; im elektronischen Antrag ist zugleich ein Beschäftigungsbeginn zum 1. April eingetragen. Ferner bitten wir um Angabe, seit welchem Datum die Aufgaben im Geschäftsbetrieb tatsächlich ausgeübt werden und ob diese im Juni bereits selbständig wahrgenommen wurden.

Diese Anforderung ist noch keine Entscheidung über Ihren Antrag. Die von Ihnen gewünschte Wirkung ab April wird anhand der vollständigen Unterlagen geprüft. Bitte übermitteln Sie keine Daten einzelner Kunden. Eine zusammenfassende Beschreibung mit anonymisierten Beispielen reicht für die Tätigkeitsdarstellung aus.

Mit freundlichen Grüßen
Sachbearbeitung Befreiung''')
add(S,18,'Versorgungswerk_Beitrag_2026','pdf','Änderung des laufenden Monatsbeitrags','Versorgungswerk der Rechtsanwältinnen und Rechtsanwälte in Hamburg','21. August 2026','Mila Ahrens, Mitgliedsnummer 48192', '''Sehr geehrte Frau Ahrens,

aufgrund der Arbeitgebermeldung vom 12. August 2026 setzen wir den laufenden Beitrag aus Ihrer Beschäftigung ab dem 1. Juli 2026 auf monatlich 1.488,00 Euro fest. Der Meldung liegt ein monatliches Bruttoarbeitsentgelt von 8.000,00 Euro zugrunde. Für Januar bis Juni wurden monatlich 1.209,00 Euro gemeldet. Die im Juli zunächst nur mit 1.209,00 Euro eingegangene Zahlung wurde durch eine Nachzahlung von 279,00 Euro am 20. August ergänzt.

Die Beitragsaufstellung erfasst die gemeldete Beschäftigung. Ihre Einkünfte aus selbständiger anwaltlicher Tätigkeit werden nach Vorlage der hierfür vorgesehenen Einkommensnachweise gesondert geprüft. Übermitteln Sie bitte nach Zugang den Einkommensteuerbescheid für 2025. Eine einmalige Einnahmenaufstellung ersetzt den Bescheid nicht, kann jedoch zur vorläufigen Zuordnung dienen.

Die Zuordnung eingehender Zahlungen zu Ihrem Mitgliedskonto enthält keine Entscheidung der gesetzlichen Rentenversicherung über eine Befreiung. Die von der Arbeitgeberin verwendete Bezeichnung des Beitrags darf deshalb nicht als ein solcher Bescheid verstanden werden.

Gegen diese Beitragsfestsetzung kann innerhalb eines Monats nach Bekanntgabe schriftlich oder in zulässiger elektronischer Form Widerspruch erhoben werden. Bitte nennen Sie dabei die Mitgliedsnummer und den Monat, dessen Beitrag Sie beanstanden.

Hildegard Ulmer, Beitragsverwaltung''')
add(S,20,'Abrechnung_Juli','pdf','Entgeltabrechnung Juli 2026','Fleetbogen Netze GmbH','31. Juli 2026','Mila Ahrens, Personalnummer 0047', '''Abrechnungszeitraum 1. Juli bis 31. Juli 2026
Eintrittsdatum in der Personalakte: 1. April 2019
Funktion: Leiterin Recht und Geschäftsbetrieb

Monatsgehalt brutto: 8.000,00 Euro.
Steuerpflichtiges Brutto: 8.000,00 Euro.
Lohnsteuer: 1.755,20 Euro. Solidaritätszuschlag: 0,00 Euro. Kirchensteuer: 0,00 Euro.
Arbeitnehmeranteil Krankenversicherung: 517,31 Euro.
Arbeitnehmeranteil Pflegeversicherung: 104,63 Euro.
Arbeitnehmeranteil Arbeitslosenversicherung: 104,00 Euro.
Einbehalt berufsständische Versorgung: 744,00 Euro.
Auszahlungsbetrag: 4.774,86 Euro.

Arbeitgeberzuschuss zur berufsständischen Versorgung: 744,00 Euro. Die für den Monat zunächst an das Versorgungswerk übermittelte Sammelzahlung beruhte noch auf dem Gehaltswert des Vormonats. Der Unterschied wird mit dem nächsten Zahlungslauf ausgeglichen. Die Differenz betrifft nicht den oben ausgewiesenen Einbehalt bei der Arbeitnehmerin.

Krankenversicherung: Nordhafen Krankenkasse; Zusatzbeitrag laut Stammdaten 3,2 Prozent. Pflegeversicherung: Elterneigenschaft hinterlegt. Abrechnung nach den bei Erstellung gespeicherten Daten. Rückfragen richten Sie an Ottilie Kröger, entgelt@fleetbogen.example.

Diese Abrechnung wurde am 31. Juli um 14:06 Uhr erstellt. Die Zahlung an Frau Ahrens wurde unter dem Verwendungszweck Gehalt 07 2026 freigegeben.''')
add(S,21,'Krankenkasse_Einkommen','eml','Unterlagen zur freiwilligen Mitgliedschaft','Nordhafen Krankenkasse <mitglieder@nordhafen-kasse.example>','2026-09-21T13:05:00+02:00','Mila Ahrens <mila.ahrens@postfach.example>', '''Sehr geehrte Frau Ahrens,

wir nehmen Bezug auf Ihre Mitteilung über die neue Arbeitgeberbezeichnung. Für die Fortführung Ihrer Versicherungsunterlagen benötigen wir die aktuelle Entgeltbescheinigung und den Einkommensteuerbescheid 2025, sobald er vorliegt. Bitte teilen Sie mit, ob die selbständige anwaltliche Tätigkeit nach Umfang oder Gewinn wesentlich verändert wurde. Angaben über private Mandatsinhalte sind hierfür nicht erforderlich.

Sie werden bei uns derzeit als freiwilliges Mitglied geführt. Ihre soziale Pflegeversicherung wird über unsere Pflegekasse verwaltet. In unserer Akte ist ein Kind eingetragen. Ein Antrag auf Befreiung von der sozialen Pflegeversicherung ist nicht vermerkt. Falls Sie dazu eine abweichende Entscheidung besitzen, übersenden Sie diese bitte in Kopie.

Die Beitragsbemessung erfolgt auf Grundlage der nachgewiesenen Einnahmen. Der von Ihnen erwähnte Bescheid zur Rentenversicherung ist hierfür nicht allein ausschlaggebend. Unser Zusatzbeitrag beträgt seit dem 1. Januar 2026 3,2 Prozent. Über die endgültige Höhe Ihrer Beiträge für frühere Zeiträume entscheiden wir nach Eingang der noch fehlenden Nachweise.

Mit freundlichen Grüßen
Leonie Wendt, Mitgliederbetreuung''')
add(S,22,'Uebergabe_vor_Urlaub','eml','Übergabe vor meinem Urlaub','Elfriede Runge <e.runge@fleetbogen.example>','2026-07-10T16:28:00+02:00','Mila Ahrens <m.ahrens@fleetbogen.example>', '''Liebe Mila,

ich habe die vier offenen Vertragsentwürfe im Ordner Recht abgelegt. Die Frist im Verfahren Linde läuft am 24. Juli; der externe Kollege kennt sie. Den Schriftsatzentwurf prüfe bitte bis Mittwoch, damit wir rechtzeitig Rückmeldung geben. Bei den drei Kulanzfällen aus dem Kundendienst liegt keine Rechtsfrage mehr auf meinem Tisch. Juna möchte von Dir eine kaufmännische Entscheidung, ob wir trotz wirksamer Laufzeitbindung Geld erstatten.

Du hattest heute gesagt, dass diese Fälle seit dem 15. Juni ständig bei Dir landen. Ich erinnere mich an zwei Besprechungen vor dem 1. Juli, in denen Juna selbst freigegeben hat. Bitte schau in die Protokolle, bevor wir gegenüber der Personalstelle ein Datum nennen. Die neuen Freigaberechte wurden nach meiner Erinnerung erst am 6. Juli im System eingerichtet.

Für August sollten wir die juristische Vertretung neu planen. Wenn Du die Morgenrunden und die Personalgespräche selbst führst, bleiben sonst alle Vertragsprüfungen bei mir. Das war im Frühjahr anders abgesprochen. Eine feste Prozentaufteilung habe ich von Dir noch nicht bekommen.

Viele Grüße und ein ruhiges Wochenende
Elfriede''')
add(S,23,'Aufgabenprotokoll_August','docx','Aufzeichnungen über die Woche vom 17 bis 21 August','Mila Ahrens','24. August 2026','Eigene Arbeitsunterlagen', '''1 Montag und Dienstag
Am Montag dauerte die Besprechung zu Kundendienstkennzahlen von 8.30 Uhr bis 10.15 Uhr. Ich wies den Teamleitern 38 offene Erstattungsfälle zu und besprach anschließend mit Juna die Kosten der verlängerten Hotlinezeiten. Von 13 Uhr bis 16 Uhr verhandelte ich die Haftungsklauseln für den Auftrag Ostkai. Am Dienstag prüfte ich vormittags den Datenschutzanhang eines Lieferanten. Nachmittags sprach ich mit zwei Mitarbeiterinnen über die Urlaubsvertretung und gab neun Erstattungen frei.

2 Mittwoch
Ich verbrachte den Vormittag mit dem Abrechnungsteam, weil 126 Kundenrechnungen doppelt erzeugt worden waren. Die technische Ursache wurde von Friedrich Knoop geprüft. Ich entschied, welche Kundengruppen zuerst angeschrieben werden. Von 14 Uhr bis 16.30 Uhr besprach ich mit dem externen Anwalt den Entwurf im Verfahren Linde. Danach beantwortete ich bis 18 Uhr Rückfragen aus dem Vertrieb zu zwei Vertragsänderungen.

3 Donnerstag und Freitag
Am Donnerstag bereitete ich eine Vertragsfreigabe für den Beirat vor und formulierte die rechtlichen Risiken in einer zweiseitigen Notiz. Juna bat mich, zusätzlich drei Varianten für die Servicekosten zu rechnen. Am Freitag nahm ich zwei Stunden an Mitarbeitergesprächen teil. Die endgültigen Gehaltserhöhungen entschied Juna. Die übrige Zeit entfiel auf das Wochenreporting, die Überarbeitung einer Kündigungsklausel und die Abstimmung der Augustabrechnung.

4 Entstehung der Aufzeichnung
Ich habe diese Notiz nach Kalender, E-Mails und Erinnerung erstellt. Pausen und kurze Unterbrechungen sind nicht erfasst. Die Notiz ist keine vollständige Stundenliste. In Kalenderterminen mit Vertrieb und Kundendienst gehen Rechtsberatung und kaufmännische Abstimmung teilweise ineinander über. Die Personalstelle hatte mich um Beispiele gebeten; eine Bewertung habe ich dieser Notiz nicht beigefügt.

Mila Ahrens''')
add(S,24,'Chat_Abrechnung','txt','Chatexport der Entgeltgruppe','Mila Ahrens, Traugott Brehm und Ottilie Kröger','4. bis 9. September 2026','Export durch Mila Ahrens', '''04.09.2026 09:41 Traugott Brehm: Der Abrechner findet nur den Bescheid mit Mobilität. Wer hat damals die neue Firma gemeldet?
04.09.2026 09:46 Mila Ahrens: Ich habe am 2. April an die Kammer geschrieben. Die Antwort steht im Ordner Zulassung.
04.09.2026 09:48 Ottilie Kröger: Den habe ich. Das ist die Kammer, nicht DRV. In unserem Programm läuft seit April dieselbe Beitragsgruppe weiter.
04.09.2026 09:55 Traugott Brehm: Bitte heute nichts rückwirkend ändern. Wir sammeln erst die Unterlagen.
07.09.2026 14:02 Mila Ahrens: Gibt es irgendwo eine Zusage des Dienstleisters, dass mein Bescheid weiter gilt?
07.09.2026 14:11 Ottilie Kröger: Nur deine Mail und Traugotts Antwort vom März. Ich habe keinen Vermerk über einen Anruf bei DRV.
07.09.2026 14:16 Mila Ahrens: Dann stelle ich den Antrag vorsorglich jetzt. Im Portal soll ich Beginn der Beschäftigung angeben. 2019 oder April?
07.09.2026 14:18 Traugott Brehm: Bitte anhand der Unterlagen selbst klären. Ich möchte da keine falsche Zahl vorgeben.
09.09.2026 08:24 Ottilie Kröger: Der Juliunterschied beim Versorgungswerk ist am 20. August bezahlt worden. Ich schicke die Zahlungsliste. Die Gehaltsabrechnung war bereits richtig, nur die Bankdatei hatte den alten Monatsbetrag.
09.09.2026 08:30 Mila Ahrens: Danke. Ich reiche auch den Julinachtrag ein. Die Aprilentscheidung erwähnt ihn noch nicht.''')
add(S,25,'Kanzlei_Nebentaetigkeit','docx','Vereinbarung zur eigenen anwaltlichen Nebentätigkeit','Fleetbogen Mobilität GmbH und Mila Ahrens','18. März 2019','Ausfertigung für beide Vertragsparteien', '''1 Gestattung
Die Fleetbogen Mobilität GmbH gestattet Frau Mila Ahrens, außerhalb ihrer vereinbarten Arbeitszeit eine eigene anwaltliche Kanzlei zu führen. Die Kanzlei wird unter ihrem Namen und auf eigene Rechnung betrieben. Frau Ahrens entscheidet selbst, welche Mandate sie annimmt. Die Gesellschaft schuldet ihr hierfür weder Vergütung noch Sachmittel.

2 Trennung der Tätigkeiten
Frau Ahrens verwendet für die Kanzlei eigene Kommunikationsmittel, Aktenführung und eine eigene Berufshaftpflichtversicherung. Kunden und Lieferanten der Gesellschaft werden nur übernommen, wenn berufsrechtlich kein Konflikt besteht und die Geschäftsführung vorab über die mögliche Berührung betrieblicher Interessen informiert wurde. Die Gesellschaft erhält keine Einsicht in fremde Mandantenakten.

3 Zeit und Nutzung von Einrichtungen
Die Nebentätigkeit darf die Erfüllung der Aufgaben aus dem Arbeitsvertrag nicht beeinträchtigen. Termine während der vertraglichen Arbeitszeit bedürfen vorheriger Freistellung oder werden im Rahmen zulässiger Gleitzeit nachgeholt. Eine pauschale Bereitschaft der Gesellschaft, Kanzleiaufwand zu vergüten, besteht nicht. Die Räume in der Werftstraße dürfen ohne besondere Vereinbarung nicht als Kanzleisitz angegeben werden.

4 Änderungen
Ändern sich Umfang oder Art der Nebentätigkeit so, dass betriebliche Interessen betroffen sein können, informiert Frau Ahrens die Gesellschaft. Eine Einschränkung der Gestattung wird nicht ohne sachlichen Grund ausgesprochen. Die Beteiligten besprechen zunächst, ob eine konkrete Interessenkollision oder zeitliche Überschneidung durch eine engere Abgrenzung behoben werden kann.

Hamburg, 18. März 2019
Traugott Brehm
Mila Ahrens''')
add(S,27,'Nachricht_zum_Junistart','eml','Was ich im Juni schon gemacht habe','Mila Ahrens <mila.ahrens@postfach.example>','2026-09-29T07:53:00+02:00','Ottilie Semmel <kanzlei@semmel-weber.example>', '''Sehr geehrte Frau Semmel,

ich habe gestern noch einmal mit Elfriede gesprochen. Mein Satz, ich hätte den Geschäftsbetrieb schon seit dem 15. Juni geleitet, war zu pauschal. Ab diesem Tag saß ich in den Besprechungen und bereitete Entscheidungen vor. Die Freigaben unterschrieb zunächst Juna. Ab Juli erwartete sie von mir eigene Entscheidungen. Das technische Recht dazu wurde erst am 6. Juli eingerichtet; an den ersten drei Arbeitstagen schickte ich Juna die Fälle per E-Mail zur Ausführung.

Es gab kein Ende meines alten Arbeitsverhältnisses im März, keine Abfindung und keine neue Probezeit. Ich habe auch keinen Neuvertrag unterschrieben. Im Personalprogramm wird dennoch der 1. April als Eintritt in Netze angezeigt. Die Übernahmevereinbarung ist die einzige Vereinbarung zu diesem Zeitpunkt.

Zu meiner Kanzlei: Die Einnahmenliste enthält einen Betrag von 1.200 Euro, der am 3. September einging, aber eine Rechnung aus Juli betrifft. Ich hatte ihn in einer frühen Aufstellung versehentlich als Julizahlung eingeordnet. Die beigefügte CSV ist nach tatsächlichem Zahlungseingang sortiert. Bitte verwenden Sie diese Fassung. Das Mandat gehörte einer Nachbarin und hatte mit Fleetbogen nichts zu tun.

Mit freundlichen Grüßen
Mila Ahrens''')
add(S,29,'Arbeitgeberbestaetigung_April','docx','Bestätigung über die Fortführung der Beschäftigung','Fleetbogen Netze GmbH','1. April 2026','Hanseatische Rechtsanwaltskammer', '''Sehr geehrte Damen und Herren,

wir bestätigen, dass Frau Mila Ahrens seit heute aufgrund der dreiseitigen Vereinbarung vom 16. März 2026 in einem vollständig übernommenen Arbeitsverhältnis bei uns tätig ist. Das zuvor mit der Fleetbogen Mobilität GmbH bestehende Arbeitsverhältnis wurde nicht beendet. Wir sind in sämtliche Rechte und Pflichten eingetreten. Die Betriebszugehörigkeit seit dem 1. April 2019 wird anerkannt.

Frau Ahrens bearbeitet weiterhin die in der Tätigkeitsbeschreibung vom 18. März 2019 genannten Rechtsfragen. Sie verhandelt Verträge, berät die Geschäftsführung und die Fachabteilungen und bearbeitet rechtliche Auseinandersetzungen. Ihre vertraglich vereinbarte fachliche Unabhängigkeit besteht unverändert fort. Arbeitsort und regelmäßige Arbeitszeit sind durch die Übernahme nicht geändert worden.

Eine weitergehende organisatorische Neuordnung wird intern besprochen. Hierzu ist bislang keine Vertragsänderung abgeschlossen worden. Sollte Frau Ahrens künftig zusätzliche Aufgaben übernehmen, wird die tatsächliche Ausgestaltung gesondert beschrieben. Die heutige Bestätigung bezieht sich allein auf die zum Übernahmestichtag fortgeführte Tätigkeit.

Für Rückfragen zu den Vertragsunterlagen steht die Geschäftsführung zur Verfügung. Diese Bestätigung wird Frau Ahrens zur Vorlage im Zusammenhang mit ihrer bisherigen Zulassung übergeben.

Hamburg, 1. April 2026
Juna Martens, Geschäftsführerin''')
add(S,28,'Arbeitgeberbestaetigung_September','docx','Beschreibung der derzeitigen Aufgaben von Frau Ahrens','Fleetbogen Netze GmbH','29. September 2026','Hanseatische Rechtsanwaltskammer und Mila Ahrens', '''1 Vertraglicher Verlauf
Das Arbeitsverhältnis von Frau Mila Ahrens wurde zum 1. April 2026 durch dreiseitigen Vertrag mit allen Rechten und Pflichten übernommen. Bis zum 30. Juni galt die bisherige Aufgabenbeschreibung. Eine Beendigung und Neubegründung wurde nicht vereinbart. Die Arbeitszeit beträgt unverändert 38 Wochenstunden.

2 Aufgaben seit Juli
Frau Ahrens bearbeitet Vertragsfragen, verhandelt rechtliche Klauseln und berät die Geschäftsführung. Daneben koordiniert sie Kundendienst und Abrechnung, führt die wöchentlichen Teamleiterrunden und entscheidet über Erstattungen bis 15.000 Euro im Einzelfall. Für operative Entscheidungen erhält sie Vorgaben zu Budget und Kundenservice. Über Einstellungen und Kündigungen entscheidet die Geschäftsführerin. Frau Ahrens ist kein Organ der Gesellschaft und verfügt nicht über Prokura.

3 Zeitanteile und Unabhängigkeit
Die Rechtsberatung erfolgt fachlich unabhängig. Für die kaufmännischen Aufgaben gelten die Vorgaben der Geschäftsführung. Eine belastbare Stundenaufteilung können wir mangels entsprechender Erfassung nicht bescheinigen. Im August entfiel nach Einschätzung der Geschäftsführerin ungefähr die Hälfte der Tätigkeit auf rechtliche Aufgaben. Die ursprüngliche Planung hatte für Recht einen höheren Anteil vorgesehen. Der tatsächliche Bedarf schwankt nach Projektphase.

4 Übergangsphase
Die Einarbeitung begann am 15. Juni. Vor dem 1. Juli traf die Geschäftsführerin die abschließenden operativen Freigaben. Ab Juli wurden diese Frau Ahrens innerhalb des vereinbarten Rahmens übertragen. Die technischen Systemrechte wurden am 6. Juli freigeschaltet. Diese Verzögerung sollte den vertraglichen Beginn nicht verschieben.

Hamburg, 29. September 2026
Juna Martens, Geschäftsführerin''')

# Zweite Akte: Organfunktion, andere Gesellschaft und gesonderter Auftrag.
add(A,1,'Mandatsanfrage_Brandt','eml','Drei Tätigkeiten und eine Sammelabrechnung','Dr. Elif Brandt <elif.brandt@postfach.example>','2026-09-28T20:17:00+02:00','Friedrich Albers <kanzlei@albers-semmel.example>', '''Sehr geehrter Herr Albers,

seit April gehöre ich dem Vorstand der Leinefaden Systeme AG an. Daneben sitze ich im Aufsichtsrat der eigenständigen Havelstrom Speicher AG und habe für die Mohnwinkel Service GmbH ein begrenztes Projekt zur Lieferantenbewertung übernommen. Unsere Entgeltstelle hat mir gesagt, ein AG-Vorstand sei generell aus der Sozialversicherung heraus. Meine Krankenkasse möchte jetzt trotzdem Unterlagen. Ich möchte die drei Tätigkeiten sauber prüfen lassen, bevor irgendjemand die Meldungen rückwirkend verändert.

Die Gesellschaften arbeiten zusammen, sind aber nach meiner Kenntnis nicht sämtlich in einem Konzern. Leinefaden hält 30 Prozent an Mohnwinkel; Havelstrom hat ganz andere Gesellschafter. Ich selbst besitze keine Aktien und keine GmbH-Anteile. Vor der Bestellung war ich bei Leinefaden angestellte Finanzleiterin. Die damalige Vereinbarung wurde Ende März beendet.

Meine Vergütung als Vorstandsmitglied ist niedriger als das frühere Festgehalt; dafür wurde eine variable Zahlung besprochen. Der Aufsichtsrat hat hierzu noch keine Zielvereinbarung beschlossen. Für die Beratung habe ich zwei Rechnungen geschrieben. Einen eigenen Mitarbeiter beschäftige ich nicht. Der Projektleiter hat mich zuletzt wie eine interne Kollegin in die Dienstagsrunde eingeladen.

Bitte berücksichtigen Sie auch meine Kranken- und Pflegeversicherung. Ich bin seit 2022 freiwillig gesetzlich versichert und habe zwei Kinder, geboren 2008 und 2012. Die Unterlagen erhalten Sie über den Ordner. Es soll zunächst keine Nachricht an die Gesellschaften versandt werden.

Mit freundlichen Grüßen
Elif Brandt''')
add(A,2,'Bestellung_Vorstand','docx','Beschluss über die Bestellung von Dr Elif Brandt','Aufsichtsrat der Leinefaden Systeme AG','18. März 2026','Dr. Elif Brandt und Vorstand der Gesellschaft', '''1 Sitzung und Teilnahme
Der Aufsichtsrat der Leinefaden Systeme AG trat am 18. März 2026 um 10 Uhr am Gesellschaftssitz Kanalhof 7, 30159 Hannover zusammen. Teilgenommen haben die Vorsitzende Ottilie Rabenau, der stellvertretende Vorsitzende Nils Berger und Mitglied Traugott Ahlers. Alle Mitglieder bestätigten den rechtzeitigen Zugang der Einladung und die Beschlussfähigkeit.

2 Bestellung
Der Aufsichtsrat bestellt Frau Dr. Elif Brandt mit Wirkung zum 1. April 2026 bis zum Ablauf des 31. März 2029 zum Mitglied des Vorstands. Sie übernimmt das Ressort Finanzen und Organisation. Die Bestellung wird einstimmig beschlossen. Frau Brandt erklärt in der Sitzung, dass sie die Bestellung annimmt. Eine Beteiligung am Grundkapital ist damit nicht verbunden.

3 Dienstvertrag und Vertretung
Die Vorsitzende wird ermächtigt, den in der Sitzung besprochenen Dienstvertrag zu unterzeichnen. Die monatliche Festvergütung soll 6.200 Euro betragen. Über eine variable Vergütung entscheidet der Aufsichtsrat anhand einer noch zu beschließenden jährlichen Zielvereinbarung. Der Abschluss des Dienstvertrags ersetzt die gesondert beschlossene Bestellung nicht. Die gesellschaftsrechtliche Vertretung der Gesellschaft richtet sich nach Satzung und Registereintragung.

4 Bisheriges Arbeitsverhältnis
Das seit 2022 bestehende Arbeitsverhältnis als Finanzleiterin soll zum 31. März 2026 durch eine gesonderte Vereinbarung beendet werden. Eine automatische Wiederaufnahme nach Ende des Vorstandsmandats wird nicht zugesagt. Die offenen Urlaubs- und Vergütungsansprüche sind in der Beendigungsvereinbarung abzurechnen.

5 Vollzug
Der Vorstand wird gebeten, die Anmeldung zum Handelsregister in Abstimmung mit dem Notariat zu veranlassen. Die Personalstelle erhält eine Abschrift für die Abrechnung. Der Beschluss enthält keine Feststellung der Versicherungspflicht einzelner Vergütungsbestandteile.

Hannover, 18. März 2026
Ottilie Rabenau, Vorsitzende
Nils Berger
Traugott Ahlers''')
add(A,3,'Vorstandsdienstvertrag','docx','Dienstvertrag für das Vorstandsmitglied Finanzen','Leinefaden Systeme AG und Dr. Elif Brandt','20. März 2026','Ausfertigung für beide Vertragsparteien', '''Zwischen der Leinefaden Systeme AG, Kanalhof 7, 30159 Hannover, vertreten durch die Aufsichtsratsvorsitzende Ottilie Rabenau, und Frau Dr. Elif Brandt, Lindensteg 23, 30167 Hannover, wird folgender Dienstvertrag geschlossen.

1 Beginn und Dauer
Das Dienstverhältnis beginnt am 1. April 2026 und endet mit Ablauf des 31. März 2029, ohne dass es einer Kündigung bedarf. Die gesellschaftsrechtliche Bestellung richtet sich nach dem Aufsichtsratsbeschluss vom 18. März 2026. Ein vorzeitiger Widerruf der Bestellung beendet den Dienstvertrag nicht ohne Weiteres. Das Recht beider Parteien zur Kündigung aus wichtigem Grund bleibt unberührt.

2 Aufgaben und Verantwortung
Frau Brandt führt das Ressort Finanzen und Organisation in eigener Verantwortung im Rahmen der gesetzlichen Pflichten, der Satzung und der Geschäftsordnung. Sie erstellt die Finanzplanung, verantwortet die Zahlungsfähigkeit und organisiert das Rechnungswesen. Die Gesamtverantwortung des Vorstands bleibt bestehen. Der Aufsichtsrat überwacht die Tätigkeit und erhält die gesetzlich und vertraglich vorgesehenen Berichte. Operative Einzelweisungen aus einem Arbeitsverhältnis werden mit diesem Vertrag nicht vereinbart.

3 Zusammenarbeit und Zeit
Frau Brandt widmet der Gesellschaft die zur ordnungsgemäßen Aufgabenerfüllung erforderliche Arbeitskraft. Ein monatliches Stundenkontingent wird nicht festgelegt. Termine außerhalb der üblichen Bürozeiten können erforderlich sein. Die Gesellschaft stellt einen Arbeitsplatz und die notwendigen Kommunikationsmittel zur Verfügung. Die Ressortverteilung und die Vertretung während Abwesenheiten werden durch den Gesamtvorstand abgestimmt.

4 Festvergütung
Frau Brandt erhält monatlich 6.200 Euro brutto, zahlbar am letzten Bankarbeitstag. Die Gesellschaft führt die gesetzlich geschuldeten Steuern und Beiträge ab und zahlt gesetzlich geschuldete Versicherungszuschüsse. Diese Klausel enthält keine eigenständige Zusage einer Versicherungsfreiheit. Nachgewiesene dienstlich veranlasste Reisekosten werden nach der Reisekostenregelung der Gesellschaft erstattet.

5 Variable Vergütung
Eine variable Jahresvergütung kann bis zu 24.000 Euro betragen. Ihre Gewährung setzt eine schriftlich beschlossene Zielvereinbarung des Aufsichtsrats und die spätere Feststellung der Zielerreichung voraus. Für 2026 ist die Zielvereinbarung bis zum 30. Juni vorgesehen. Aus der bloßen Nennung des Höchstbetrags folgt weder eine garantierte Mindestzahlung noch eine anteilige monatliche Fälligkeit. Kommt eine Vereinbarung nicht zustande, beraten die Parteien über die entstandene Situation, ohne mit dieser Regelung einen bestimmten Zahlungsbetrag vorwegzunehmen.

6 Urlaub und Verhinderung
Frau Brandt stimmt Erholungszeiten mit dem Vorstandsvorsitzenden ab und sorgt für die notwendige Vertretung. Als Richtgröße werden 30 Arbeitstage pro Kalenderjahr vereinbart. Bei unverschuldeter Dienstunfähigkeit wird die Festvergütung bis zu sechs Monate fortgezahlt, soweit nicht anderweitige Leistungen anzurechnen sind. Längere Verhinderungen werden der Aufsichtsratsvorsitzenden unverzüglich mitgeteilt.

7 Nebentätigkeiten
Das Aufsichtsratsmandat bei der Havelstrom Speicher AG und das zeitlich begrenzte Beratungsprojekt für die Mohnwinkel Service GmbH werden nach Maßgabe des gesonderten Zustimmungsbeschlusses gestattet. Frau Brandt zeigt Interessenkonflikte an und nimmt an betroffenen Entscheidungen nicht ohne vorherige Abstimmung teil. Die Zustimmung übernimmt keine Abrechnung oder sozialversicherungsrechtliche Beurteilung des jeweiligen Rechtsverhältnisses.

8 Vertraulichkeit und Rückgabe
Vertrauliche Geschäftsunterlagen werden nur zur Aufgabenerfüllung verwendet und bei Ende des Dienstverhältnisses zurückgegeben. Die Verschwiegenheitspflicht gilt auch danach. Gesetzliche Auskunfts-, Anzeige- und Mitwirkungspflichten werden nicht eingeschränkt. Personenbezogene Daten werden nach den dafür geltenden Regeln verarbeitet.

9 Schlussbestimmungen
Eine Rückkehr in das frühere Arbeitsverhältnis ist nicht vereinbart. Änderungen werden schriftlich dokumentiert; zwingende Formvorschriften und wirksame individuelle Abreden bleiben unberührt. Sollte eine einzelne Vereinbarung unwirksam sein, wird eine zulässige Regelung vereinbart, die dem beabsichtigten wirtschaftlichen Zweck möglichst nahekommt, ohne zwingendes Recht zu umgehen.

Hannover, 20. März 2026
Ottilie Rabenau für den Aufsichtsrat
Dr. Elif Brandt''')
add(A,4,'Registermitteilung','pdf','Mitteilung über die Eintragung der Vorstandsänderung','Notariat Dr. Walther Finke','9. April 2026','Leinefaden Systeme AG, Aufsichtsratsvorsitzende Ottilie Rabenau', '''Vorgang 126 2026

Sehr geehrte Frau Rabenau,

das Registergericht hat die am 24. März angemeldete Vorstandsänderung am 8. April 2026 eingetragen. Dr. Elif Brandt ist seit dem 1. April 2026 zum Vorstandsmitglied bestellt. Nach der mitgeteilten Eintragung wird die Gesellschaft durch zwei Vorstandsmitglieder gemeinsam oder durch ein Vorstandsmitglied zusammen mit einem Prokuristen vertreten. Eine Einzelvertretungsbefugnis für Frau Brandt wurde nicht angemeldet.

Die Eintragung betrifft die Leinefaden Systeme AG mit Sitz in Hannover. Der vorhandene Vorstandsvorsitzende bleibt im Amt. Die Gesellschaft hat uns keine Änderung des Grundkapitals und keinen Erwerb von Aktien durch Frau Brandt zur Beurkundung vorgelegt. Die Anmeldung beruhte auf dem Aufsichtsratsbeschluss vom 18. März und der Annahmeerklärung der Bestellten.

Bitte legen Sie die Registermitteilung zur gesellschaftsrechtlichen Akte. Für Fragen des Dienstvertrags und seiner Abrechnung ist der unterzeichnete Vertrag maßgeblich hinzuzunehmen. Diese Mitteilung enthält keine Auskunft über Versicherungs- oder Beitragspflichten.

Die gesonderte Ausfertigung des Registerauszugs wird der Geschäftsanschrift postalisch übersandt. Für Rückfragen zu diesem Vollzug verwenden Sie bitte die oben genannte Vorgangsnummer.

Mit freundlichen Grüßen
Dr. Walther Finke, Notar''')
add(A,5,'Satzung_Auszug','docx','Satzungsbestimmungen über Vorstand und Aufsichtsrat','Leinefaden Systeme AG','Fassung vom 12. November 2025','Auszug aus der Gesellschaftsakte', '''1 Firma und Sitz
Die Gesellschaft führt die Firma Leinefaden Systeme AG und hat ihren Sitz in Hannover. Gegenstand des Unternehmens ist die Entwicklung und der Vertrieb technischer Systeme für stationäre Energiespeicherung sowie damit verbundener Dienstleistungen. Das Geschäftsjahr ist das Kalenderjahr.

2 Vorstand
Der Vorstand besteht aus mindestens zwei Personen. Die Zahl seiner Mitglieder bestimmt der Aufsichtsrat. Der Aufsichtsrat kann ein Mitglied zum Vorsitzenden ernennen. Der Vorstand leitet die Gesellschaft unter eigener Verantwortung. Er gibt sich eine Geschäftsordnung, soweit nicht der Aufsichtsrat eine solche erlässt. Eine Ressortverteilung entbindet kein Mitglied von seiner gesetzlichen Gesamtverantwortung.

3 Vertretung
Die Gesellschaft wird durch zwei Vorstandsmitglieder gemeinsam oder durch ein Vorstandsmitglied zusammen mit einem Prokuristen vertreten. Der Aufsichtsrat kann einzelnen Mitgliedern Einzelvertretungsbefugnis erteilen, soweit dies rechtlich zulässig ist. Interne Zustimmungsvorbehalte beschränken die Vertretungsmacht gegenüber Dritten nicht über die gesetzlichen Grenzen hinaus.

4 Zustimmungspflichtige Geschäfte
Der Erwerb oder die Veräußerung von Beteiligungen, die Aufnahme von Darlehen über 500.000 Euro und der Abschluss von Verträgen außerhalb des genehmigten Budgets mit einem Gesamtwert über 250.000 Euro bedürfen der vorherigen Zustimmung des Aufsichtsrats. Der Aufsichtsrat kann für einzelne Geschäftsarten weitere Zustimmungsvorbehalte beschließen. Die laufende Geschäftsführung wird dadurch nicht auf ihn übertragen.

5 Aufsichtsrat
Der Aufsichtsrat besteht aus drei Mitgliedern, soweit zwingendes Recht keine andere Zusammensetzung verlangt. Er wählt aus seiner Mitte eine Vorsitzende und einen Stellvertreter. Beschlüsse werden mit einfacher Stimmenmehrheit gefasst, sofern das Gesetz oder diese Satzung keine andere Mehrheit verlangt. Über Sitzungen wird eine von der Vorsitzenden unterzeichnete Niederschrift erstellt.

6 Vergütung und Berichte
Die Vergütung des Vorstands wird durch den Aufsichtsrat vereinbart. Der Vorstand berichtet dem Aufsichtsrat über Planung, Geschäftsentwicklung, Liquidität und wesentliche Risiken. Die Vergütung des Aufsichtsrats richtet sich nach dem Beschluss der Hauptversammlung. Eine Vergütung für gesonderte Dienstleistungen bedarf einer eigenständigen Vereinbarung und der erforderlichen gesellschaftsrechtlichen Zustimmung.

Abschrift erstellt durch Ottilie Rabenau am 22. September 2026. Die übrigen Satzungsbestimmungen sind in diesem Auszug nicht wiedergegeben.''')
add(A,6,'Arbeitsvertrag_Finanzleitung_2022','docx','Arbeitsvertrag als Leiterin Finanzen','Leinefaden Systeme AG und Dr. Elif Brandt','14. Dezember 2021','Ausfertigung für Dr. Elif Brandt', '''1 Beginn und Aufgaben
Frau Dr. Elif Brandt wird ab dem 1. Januar 2022 als Leiterin Finanzen beschäftigt. Sie untersteht dem Vorstandsvorsitzenden und setzt dessen Vorgaben zur Budgetplanung und Organisation des Rechnungswesens um. Sie bereitet Investitionsentscheidungen vor, ohne selbst Mitglied eines Gesellschaftsorgans zu sein. Über Budgetabweichungen von mehr als 25.000 Euro entscheidet der zuständige Vorstand.

2 Arbeitszeit und Arbeitsort
Die regelmäßige Arbeitszeit beträgt 40 Wochenstunden an fünf Arbeitstagen. Arbeitsort ist Hannover. Mobile Arbeit an bis zu zwei Tagen je Woche wird nach Abstimmung gestattet. Angeordnete Mehrarbeit wird durch Freizeit ausgeglichen. Dienstreisen sind im für die Aufgabe üblichen Umfang durchzuführen. Die Gesellschaft stellt die erforderlichen Geräte und Zugänge bereit.

3 Vergütung
Die feste monatliche Bruttovergütung beträgt 7.200 Euro und wird zum Monatsende gezahlt. Ein variabler Vergütungsanspruch wird nicht vereinbart. Steuern und gesetzlich geschuldete Beiträge werden abgeführt. Gesetzliche Zuschüsse zur Kranken- und Pflegeversicherung werden nach Vorlage der Versicherungsnachweise gezahlt. Notwendige Reisekosten werden gegen Beleg erstattet.

4 Urlaub und Krankheit
Der jährliche Urlaub beträgt bei einer Fünftagewoche 30 Arbeitstage. Die Lage des Urlaubs wird unter Berücksichtigung der betrieblichen Belange abgestimmt. Arbeitsunfähigkeit und voraussichtliche Dauer sind unverzüglich mitzuteilen. Die Entgeltfortzahlung richtet sich nach dem Gesetz. Die gesetzlichen Anzeige- und Nachweispflichten gelten in ihrer jeweiligen Fassung.

5 Nebentätigkeit und Vertraulichkeit
Entgeltliche Nebentätigkeiten sind vor Aufnahme anzuzeigen. Eine Untersagung erfolgt nur bei berechtigten betrieblichen Interessen. Geschäftsgeheimnisse und personenbezogene Informationen werden vertraulich behandelt. Unterlagen und Geräte sind bei Vertragsende zurückzugeben; gesetzlich zulässige eigene Nachweise bleiben unberührt.

6 Beendigung
Nach einer sechsmonatigen Probezeit gilt eine Kündigungsfrist von drei Monaten zum Quartalsende. Zwingende längere Fristen bleiben maßgeblich. Kündigungen bedürfen der Schriftform. Eine Bestellung in ein Gesellschaftsorgan verändert dieses Arbeitsverhältnis nur aufgrund einer ausdrücklichen weiteren Vereinbarung.

Hannover, 14. Dezember 2021
Lennart Volkmann für die Gesellschaft
Dr. Elif Brandt''')
add(A,7,'Beendigung_Arbeitsverhaeltnis','docx','Vereinbarung über die Beendigung des Arbeitsverhältnisses','Leinefaden Systeme AG und Dr. Elif Brandt','20. März 2026','Ausfertigung für beide Vertragsparteien', '''1 Beendigung
Die Parteien beenden das Arbeitsverhältnis aus dem Vertrag vom 14. Dezember 2021 einvernehmlich mit Ablauf des 31. März 2026. Es wird für die Dauer eines künftigen Vorstandsdienstverhältnisses nicht ruhend gestellt. Ein Anspruch auf spätere Wiederaufnahme oder Wiedereinstellung wird nicht vereinbart.

2 Abrechnung
Die feste Vergütung wird bis einschließlich März unverändert abgerechnet. Eine Abfindung wird nicht gezahlt. Aus der Reisekostenabrechnung vom 12. März sind noch 286,40 Euro zu erstatten; dieser Betrag wird mit der Märzabrechnung ausgezahlt. Die Parteien verzichten nicht auf Ansprüche, die aus einer späteren rechnerischen Korrektur dieser Abrechnung entstehen.

3 Urlaub und Arbeitsmittel
Die bis zum Beendigungsdatum entstandenen Urlaubsansprüche sind nach der Personalaufstellung durch die bereits gewährten Urlaubstage erfüllt. Frau Brandt bestätigt nicht die Erfüllung unbekannter weiterer Ansprüche. Die bislang verwendeten Geräte verbleiben für das anschließende Dienstverhältnis bei ihr. Ihre weitere Nutzung beruht ab dem 1. April auf dem neuen Dienstvertrag und nicht auf dem beendeten Arbeitsvertrag.

4 Zeugnis und Unterlagen
Die Gesellschaft stellt auf Wunsch ein qualifiziertes Zeugnis für die Tätigkeit als Finanzleiterin aus. Die gesetzlich erforderlichen Meldungen und Bescheinigungen werden erstellt. Diese Vereinbarung enthält keine verbindliche Festlegung der versicherungsrechtlichen Behandlung des anschließenden Rechtsverhältnisses.

5 Abschluss
Die Parteien haben den Text vor Unterzeichnung erhalten und besprochen. Weitere Abreden zur Beendigung bestehen nicht. Die Unterzeichnung erfolgt eigenhändig in zwei gleichlautenden Ausfertigungen.

Hannover, 20. März 2026
Lennart Volkmann für die Gesellschaft
Dr. Elif Brandt''')
add(A,9,'Krankenkasse_Nachfrage','pdf','Anforderung von Unterlagen zum Tätigkeitswechsel','Nordhafen Krankenkasse','17. September 2026','Dr. Elif Brandt, Lindensteg 23, 30167 Hannover', '''Mitgliedsnummer N 884203

Sehr geehrte Frau Brandt,

Ihr Arbeitgeber hat uns eine Änderung Ihrer Tätigkeit zum 1. April 2026 gemeldet. In den übermittelten Unterlagen ist die Bezeichnung Vorstandsmitglied vermerkt. Zur Prüfung Ihrer Mitgliedschaft und der Beitragsbemessung benötigen wir den Dienstvertrag, den Bestellungsbeschluss und Angaben zur regelmäßigen Vergütung. Bitte legen Sie auch eine gegebenenfalls abgeschlossene Zielvereinbarung vor.

Sie werden bei uns bislang als freiwilliges Mitglied geführt. Für Januar bis März liegt ein Monatsgehalt von 7.200 Euro vor. Ab April wurde ein Festgehalt von 6.200 Euro gemeldet. Ein zusätzlicher Betrag von monatlich 2.000 Euro erscheint nur in einer Planungsübersicht und wurde bislang nicht als Zahlung nachgewiesen. Bitte teilen Sie mit, ob eine verbindliche Zusage oder eine bereits entstandene Forderung besteht.

Ferner haben Sie auf weitere Tätigkeiten hingewiesen. Bitte beschreiben Sie deren Art, Umfang und Einnahmen. Hierfür benötigen wir keine vollständigen vertraulichen Projektakten, sondern die jeweiligen Vereinbarungen und eine Einnahmenübersicht. Die Elterneigenschaft für Ihre Kinder, geboren 2008 und 2012, ist gespeichert.

Bitte reichen Sie die Unterlagen bis zum 9. Oktober 2026 ein. Über Ihren Versicherungsverlauf und gegebenenfalls geänderte Beiträge wird erst nach Prüfung entschieden. Diese Anforderung ist noch keine Festsetzung einer Nachzahlung.

Mit freundlichen Grüßen
Leonie Wendt, Mitgliederbetreuung''')
add(A,10,'Entgeltstelle_Rundmail','eml','Abrechnung Organwechsel ab April','Elfriede Holle <entgelt@leinefaden.example>','2026-04-10T10:31:00+02:00','Dr. Elif Brandt <e.brandt@leinefaden.example>', '''Guten Tag Frau Brandt,

wir haben Ihre Funktion im Personalprogramm auf Vorstand umgestellt. In der internen Liste steht dafür der Hinweis keine SV. Ich habe deshalb zunächst die Einstellungen aus dem bisherigen Datensatz nicht einzeln geprüft, sondern den Vorgang an unseren Abrechnungsdienstleister weitergegeben. Die Krankenversicherung soll nach dessen Rückfrage vorerst so laufen wie bisher.

Bitte reichen Sie die Bescheinigung Ihrer Krankenkasse zu den aktuellen Beiträgen ein. Für die Pflegeversicherung sind Ihre beiden Kinder gespeichert. Eine private Versicherung ist in unserem System nicht hinterlegt. Die Aufsichtsratsvergütung bei Havelstrom zahlen wir nicht aus und kennen deren Abrechnung nicht.

Die Beratung für Mohnwinkel soll nach Auskunft von Herrn Volkmann direkt von Ihnen in Rechnung gestellt werden. Bitte schicken Sie solche Rechnungen nicht an unsere Entgeltstelle. Mohnwinkel hat eine eigene Buchhaltung. Ob dort dieselbe Organregelung gelten soll, wurde mit uns nicht besprochen. Ich habe lediglich vermerkt, dass der Aufsichtsrat die Nebentätigkeit erlaubt hat.

Freundliche Grüße
Elfriede Holle''')
add(A,11,'Zustimmung_Nebentaetigkeiten','docx','Zustimmung zu den angezeigten Nebentätigkeiten','Aufsichtsrat der Leinefaden Systeme AG','22. März 2026','Dr. Elif Brandt', '''1 Beschlussgegenstand
Der Aufsichtsrat hat die von Frau Dr. Elif Brandt angezeigten Tätigkeiten bei der Havelstrom Speicher AG und der Mohnwinkel Service GmbH beraten. Frau Brandt hat die wesentlichen Vertragsunterlagen vorgelegt und erklärt, dass sie keine Anteile an diesen Gesellschaften hält. Die Zustimmung wird einstimmig unter den nachfolgenden Bedingungen erteilt.

2 Aufsichtsratsmandat
Die Übernahme des Aufsichtsratsmandats bei der Havelstrom Speicher AG wird gestattet. Frau Brandt informiert die Vorsitzende über absehbare Interessenkonflikte zwischen den Gesellschaften. Bei solchen Konflikten werden die erforderlichen Vorkehrungen für Teilnahme, Information und Abstimmung im konkreten Fall festgelegt. Eine Verpflichtung, operative Leistungen für Havelstrom zu erbringen, wird durch diese Zustimmung nicht begründet.

3 Beratungsprojekt
Das Projekt zur Lieferantenbewertung bei Mohnwinkel wird für den Zeitraum April bis September 2026 gestattet. Der vorgesehene Umfang beträgt höchstens 16 Beratungstage. Die Leistung wird von Frau Brandt unmittelbar gegenüber Mohnwinkel erbracht und abgerechnet. Die Leinefaden Systeme AG schuldet aus diesem Projekt keine Vergütung und übernimmt keine Gewähr für dessen wirtschaftlichen Erfolg.

4 Wahrung der Vorstandspflichten
Die Nebentätigkeiten dürfen die Erfüllung der Vorstandspflichten nicht beeinträchtigen. Eine wesentliche Ausweitung des Projekts oder eine dauerhafte Eingliederung in eine zusätzliche Funktion ist erneut anzuzeigen. Der Aufsichtsrat behält sich vor, bei einem konkret nachgewiesenen Konflikt eine Anpassung zu verlangen. Eine Aussage zur versicherungsrechtlichen Einordnung der zusätzlichen Tätigkeiten wird nicht getroffen.

Hannover, 22. März 2026
Ottilie Rabenau, Vorsitzende des Aufsichtsrats''')
add(A,12,'Havelstrom_Wahlbeschluss','docx','Wahl von Dr Elif Brandt in den Aufsichtsrat','Hauptversammlung der Havelstrom Speicher AG','26. März 2026','Abschrift für Dr. Elif Brandt', '''1 Versammlung
Die Hauptversammlung der Havelstrom Speicher AG fand am 26. März 2026 am Gesellschaftssitz in Hannover statt. Sämtliche stimmberechtigten Aktionäre waren vertreten. Den Vorsitz führte Waltraud König. Die Versammlung behandelte die Nachwahl eines Aufsichtsratsmitglieds nach dem Ausscheiden von Herrn Berger.

2 Wahl
Frau Dr. Elif Brandt wird mit Wirkung zum 1. April 2026 für die verbleibende Amtszeit bis zur Beendigung der ordentlichen Hauptversammlung 2028 in den Aufsichtsrat gewählt. Frau Brandt nimmt die Wahl an. Sie erhält weder eine Geschäftsführungsbefugnis noch eine Vollmacht zur laufenden Vertretung der Gesellschaft. Die Wahl erfolgt einstimmig.

3 Vergütung
Für das Mandat wird eine feste Jahresvergütung von 12.000 Euro festgesetzt. Beginnt oder endet das Mandat während des Kalenderjahrs, erfolgt die Vergütung zeitanteilig nach vollen Monaten. Für 2026 ergibt sich bei unverändertem Verlauf ein Betrag von 9.000 Euro. Die Zahlung ist nach Ablauf des Geschäftsjahrs fällig. Nachgewiesene notwendige Reisekosten werden zusätzlich erstattet. Ein Sitzungsgeld wird nicht gewährt.

4 Aufgaben
Frau Brandt nimmt an der Überwachung und Beratung des Vorstands teil. Sie soll insbesondere die Berichte über Beschaffung und Liquidität kritisch prüfen. Diese Zuständigkeit innerhalb des Aufsichtsrats begründet keinen gesonderten Auftrag zur Erstellung operativer Einkaufslisten oder zum Abschluss von Lieferantenverträgen. Ein solcher Auftrag müsste gesondert beraten und dokumentiert werden.

Hannover, 26. März 2026
Waltraud König, Versammlungsleiterin
Dr. Elif Brandt, Annahmeerklärung''')
add(A,13,'Havelstrom_Sitzungsprotokoll','docx','Niederschrift der Aufsichtsratssitzung vom 15 Juni','Havelstrom Speicher AG','15. Juni 2026','Aufsichtsratsmitglieder und Vorstand', '''1 Berichte des Vorstands
Der Vorstand erläuterte den Stand der Lieferantenverhandlungen und die Entwicklung des Zahlungsmittelbestands. Die geplante Lieferung von Wechselrichtern verschiebt sich nach derzeitigem Stand um vier Wochen. Frau Brandt fragte nach der Absicherung der Anzahlungen und bat um eine nach Fälligkeit gegliederte Übersicht. Der Vorstand sagte die Übersendung bis zum 22. Juni zu.

2 Beschaffung und Überwachung
Frau Brandt regte an, die Abhängigkeit von einem Lieferanten im Risikobericht gesondert darzustellen. Sie übernahm es, die im Aufsichtsrat vorhandenen Fragen zu sammeln. Ein Mitglied schlug vor, Frau Brandt solle die operative Lieferantensuche selbst führen. Die Vorsitzende stellte klar, dass dies in der Sitzung nicht als zusätzlicher Auftrag beschlossen werde. Der Vorstand blieb für Auswahl, Verhandlung und Abschluss zuständig.

3 Berichtspflichten
Bis zur nächsten Sitzung soll der Vorstand eine Monatsübersicht über Zahlungseingänge und Verpflichtungen vorlegen. Der Aufsichtsrat erhält hierzu denselben Unterlagensatz. Eine laufende Teilnahme seiner Mitglieder an den internen Montagsrunden des Vorstands ist nicht vorgesehen. Einzelne Rückfragen können schriftlich gestellt werden.

4 Vergütung und Auslagen
Die jährliche Mandatsvergütung wird nach dem Hauptversammlungsbeschluss vom 26. März abgerechnet. Frau Brandt reichte Fahrtbelege über 68,40 Euro ein. Eine gesonderte Beratungshonorierung wurde nicht beschlossen. Die Vorsitzende bittet darum, außerhalb der Sitzungen gestellte Fragen eindeutig als Aufsichtsratsanfragen zu kennzeichnen.

Waltraud König, Vorsitzende
Protokoll versandt am 18. Juni 2026''')
add(A,14,'Beratungsvertrag_Mohnwinkel','docx','Vertrag über die Bewertung des Lieferantenprozesses','Mohnwinkel Service GmbH und Dr. Elif Brandt','30. März 2026','Ausfertigung für beide Parteien', '''Zwischen der Mohnwinkel Service GmbH, Speicherweg 9, 30179 Hannover, vertreten durch Geschäftsführer Friedrich Knoop, und Dr. Elif Brandt, Lindensteg 23, 30167 Hannover, wird folgender Vertrag geschlossen.

1 Auftrag
Frau Brandt untersucht den bestehenden Prozess zur Freigabe und Bewertung von Lieferanten. Sie erstellt bis zum 30. September 2026 einen schriftlichen Bericht mit einer Bestandsaufnahme, einer Liste festgestellter Lücken und Vorschlägen zur Neuordnung. Die Umsetzung der Vorschläge und die Entscheidung über einzelne Lieferanten bleiben bei der Auftraggeberin. Eine Bestellung zur Geschäftsführerin oder Prokuristin ist nicht Gegenstand des Vertrags.

2 Durchführung
Die Auftragnehmerin bestimmt Ort und zeitliche Lage ihrer Tätigkeit grundsätzlich selbst. Sie stimmt erforderliche Gespräche mit den Ansprechpartnern ab. Die Auftraggeberin stellt die vorhandenen Prozessbeschreibungen und die erforderlichen sachlichen Informationen zur Verfügung. Soweit aus Datenschutzgründen ein gesicherter Zugriff erforderlich ist, wird ein eingeschränkter Systemzugang bereitgestellt. Eine allgemeine Verpflichtung zur täglichen Anwesenheit wird nicht vereinbart.

3 Umfang und Vergütung
Es sind höchstens 16 Beratungstage vorgesehen. Ein Beratungstag umfasst acht Stunden; kürzere Einsätze werden anteilig erfasst. Das Honorar beträgt 900 Euro je Beratungstag zuzüglich gesetzlicher Umsatzsteuer. Zusätzliche Tage bedürfen vorheriger schriftlicher Vereinbarung. Die Abrechnung erfolgt monatlich mit einer nachvollziehbaren Beschreibung der geleisteten Tätigkeiten. Rechnungen sind binnen 14 Tagen nach Zugang zu bezahlen.

4 Zusammenarbeit
Für Abstimmungen benennt die Auftraggeberin Projektleiter Nils Krause. Er koordiniert Gesprächspartner und Unterlagen. Änderungen des Leistungsgegenstands werden zwischen den Vertragsparteien vereinbart. Die Auftragnehmerin darf andere Aufträge annehmen und setzt eigene Arbeitsmittel ein, soweit der sichere Zugriff keine Geräte der Auftraggeberin verlangt. Ein Vertretungseinsatz fachlich geeigneter Personen bedarf wegen der vertraulichen Unterlagen vorheriger Zustimmung.

5 Verschwiegenheit und Daten
Die Auftragnehmerin behandelt nicht öffentliche Informationen vertraulich und verwendet sie ausschließlich für den Auftrag. Personenbezogene Daten werden nur soweit erforderlich verarbeitet. Nach Projektende werden Kopien zurückgegeben oder gelöscht, soweit keine gesetzlichen Aufbewahrungspflichten entgegenstehen. In Veröffentlichungen wird der Name der Auftraggeberin nur mit Zustimmung verwendet.

6 Bericht und Verantwortung
Der Bericht wird als PDF und bearbeitbare Datei übergeben. Die Auftraggeberin darf ihn intern nutzen und ihren Beratern zur Prüfung überlassen. Frau Brandt erläutert Rückfragen innerhalb von zwei Wochen nach Übergabe. Sie schuldet eine sorgfältige fachliche Bearbeitung, jedoch keinen bestimmten Einsparungsbetrag. Die Haftung richtet sich nach den gesetzlichen Vorschriften.

7 Beendigung
Der Vertrag endet mit vollständiger Leistungserbringung, spätestens am 30. September 2026. Beide Parteien können aus wichtigem Grund kündigen. Bereits erbrachte und nachgewiesene Leistungen werden vergütet. Bei vorzeitiger Beendigung werden die vorhandenen Arbeitsstände geordnet übergeben. Weitergehende gesetzliche Ansprüche bleiben unberührt.

Hannover, 30. März 2026
Friedrich Knoop
Dr. Elif Brandt''')
add(A,15,'Einladung_Dienstagsrunde','eml','Dienstagsrunde Einkauf ab Mai','Nils Krause <n.krause@mohnwinkel.example>','2026-04-28T14:36:00+02:00','Dr. Elif Brandt <elif.brandt@postfach.example>', '''Hallo Frau Brandt,

wir haben gesehen, dass die einzelnen Interviewtermine zu viel Abstimmung brauchen. Bitte kommen Sie ab nächster Woche dienstags um 9 Uhr in unsere Einkaufsrunde. Ich habe Sie bis Ende September in die Terminserie aufgenommen. In der Runde können Sie direkt sagen, welche Unterlagen fehlen. Der Geschäftsführer möchte außerdem, dass Sie die Freigabeliste durchsehen, bevor sie freitags an ihn geht.

Für den Zugriff auf die Lieferantendaten erhalten Sie ein Gastkonto und ein gesichertes Leihgerät. Sie dürfen aus Datenschutzgründen keine vollständigen Stammdatensätze auf Ihren privaten Rechner laden. Der Arbeitsplatz im kleinen Besprechungsraum ist für die Dienstage reserviert. Wenn Sie verhindert sind, geben Sie mir bitte vorher Bescheid, damit wir die passenden Gesprächspartner nicht unnötig einladen.

Die Zeit gehört nach meinem Verständnis zu den vereinbarten Beratungstagen. Bitte sagen Sie, wenn wir mit den wöchentlichen Terminen das Kontingent überschreiten. Die Kolleginnen behandeln Sie inzwischen wie unsere Einkaufsleitung; das ist sicher praktisch, war aber im Projektpapier anders bezeichnet. Eine zusätzliche Vollmacht habe ich für Sie nicht.

Viele Grüße
Nils Krause''')
add(A,16,'Antwort_Projektumfang','eml','Re Dienstagsrunde Einkauf ab Mai','Dr. Elif Brandt <elif.brandt@postfach.example>','2026-04-29T08:12:00+02:00','Nils Krause <n.krause@mohnwinkel.example>', '''Hallo Herr Krause,

an den ersten drei Dienstagen kann ich teilnehmen. Die Serie bis September kann ich noch nicht zusagen, weil ich für Leinefaden mehrere feste Berichtstermine habe. Ich möchte auch vermeiden, dass die Beratung unbemerkt zur laufenden Leitung Ihres Einkaufs wird. Ich kann die Liste im Hinblick auf den Freigabeprozess stichprobenweise ansehen; einzelne Lieferanten auswählen oder Bestellungen freigeben soll weiterhin Ihr Team.

Das gesicherte Gerät ist in Ordnung. Bitte richten Sie das Konto nur für die Projektordner ein. Am eigenen Bericht arbeite ich überwiegend abends in meinem Büro. Die Interviewgespräche und die Vor-Ort-Zeit halte ich in meiner Tagesliste fest. Für die zweite Maihälfte schlage ich zwei längere Termine statt mehrerer kurzer Runden vor.

Wenn Sie zusätzliche operative Unterstützung benötigen, müssen wir den Umfang und die Verantwortung gesondert besprechen. Ich habe bisher keine Mitarbeitenden dafür vorgesehen. Die 16 Tage aus unserem Vertrag sollten für die Bestandsaufnahme und den Bericht reichen, wenn ich die angeforderten Unterlagen vollständig erhalte.

Freundliche Grüße
Elif Brandt''')
add(A,17,'Rechnung_Juni','pdf','Rechnung für Beratungsleistungen im Juni','Dr. Elif Brandt','30. Juni 2026','Mohnwinkel Service GmbH, Speicherweg 9, 30179 Hannover', '''Rechnungsnummer EB 2026 06
Leistungszeitraum 1. Juni bis 30. Juni 2026

Sehr geehrter Herr Knoop,

für die Leistungen aus unserem Vertrag vom 30. März berechne ich 3,5 Beratungstage zu je 900,00 Euro. Der Nettobetrag beträgt 3.150,00 Euro. Hinzu kommen 19 Prozent Umsatzsteuer in Höhe von 598,50 Euro. Der Rechnungsbetrag beträgt 3.748,50 Euro und ist bis zum 14. Juli 2026 zu zahlen.

Die Leistung umfasste die Auswertung der Freigabeunterlagen, Gespräche mit Einkauf und Buchhaltung, eine Stichprobe aus zehn Lieferantenakten sowie die Überarbeitung der Prozessdarstellung. Zwei Dienstagsrunden wurden in das Tageskontingent einbezogen. Die Einzelzeiten ergeben sich aus der bereits an Herrn Krause übersandten Aufstellung. Für eine laufende Einkaufsleitung wird mit dieser Rechnung keine zusätzliche Pauschale berechnet.

Bitte verwenden Sie für die Zahlung das bereits in Ihrer Kreditorenakte hinterlegte Konto und geben Sie die Rechnungsnummer an. Bei Rückfragen zu einzelnen Zeitpositionen schreiben Sie bitte an elif.brandt@postfach.example. Eine weitere Beauftragung über die vereinbarten 16 Tage hinaus ist bisher nicht erfolgt.

Mit freundlichen Grüßen
Dr. Elif Brandt
Vermerk Buchhaltung Mohnwinkel: Rechnung eingegangen am 1. Juli; Zahlung freigegeben am 13. Juli 2026.''')
add(A,18,'Rechnung_August','pdf','Rechnung für Beratungsleistungen im August','Dr. Elif Brandt','31. August 2026','Mohnwinkel Service GmbH, Speicherweg 9, 30179 Hannover', '''Rechnungsnummer EB 2026 08
Leistungszeitraum 1. August bis 31. August 2026

Sehr geehrter Herr Knoop,

für die im August erbrachten Leistungen berechne ich 4 Beratungstage zu je 900,00 Euro. Der Nettobetrag beträgt 3.600,00 Euro. Die Umsatzsteuer von 19 Prozent beträgt 684,00 Euro. Der Rechnungsbetrag von 4.284,00 Euro ist bis zum 14. September 2026 zu zahlen.

Die Leistung umfasste die Nachbearbeitung fehlender Freigaben, die Auswertung weiterer Lieferantenakten und die Fertigstellung des Berichtsteils zur Trennung von Bestellung und Zahlungsfreigabe. Auf Wunsch des Projektleiters nahm ich außerdem an drei Besprechungen zu aktuellen Lieferverzögerungen teil. Soweit diese Besprechungen ausschließlich den laufenden Einkauf betrafen, habe ich die Zeit nicht doppelt als Berichtsbearbeitung erfasst.

Nach meiner bisherigen Aufstellung sind bis Ende August insgesamt 14,5 Tage verbraucht. Für die Endfassung verbleiben damit 1,5 Tage aus dem vereinbarten Höchstkontingent. Sollte die Geschäftsführung eine zusätzliche operative Begleitung wünschen, bitte ich um eine gesonderte Abstimmung vor weiterer Leistungserbringung.

Mit freundlichen Grüßen
Dr. Elif Brandt
Vermerk der Kreditorenbuchhaltung: Zahlung am 16. September 2026; verspätete Freigabe wegen Urlaub.''')
add(A,19,'Beteiligungen_Stichtag','docx','Beteiligungsübersicht zum 31 August 2026','Lennart Volkmann, Vorstand Leinefaden Systeme AG','22. September 2026','Dr. Elif Brandt', '''1 Leinefaden Systeme AG
Das Grundkapital wird nach der letzten bei der Gesellschaft hinterlegten Übersicht zu 60 Prozent von der Familie Rabenau und zu 40 Prozent von der Riedblick Beteiligungs GmbH gehalten. Dr. Elif Brandt hält keine Aktien. Sie ist seit April Mitglied des Vorstands. Eine Stimmbindungsvereinbarung mit ihr besteht nach den hier vorliegenden Unterlagen nicht.

2 Mohnwinkel Service GmbH
Die Leinefaden Systeme AG hält 30 Prozent des Stammkapitals. Weitere 45 Prozent hält Friedrich Knoop, 25 Prozent hält die Riedblick Beteiligungs GmbH. Die Gesellschafter entscheiden grundsätzlich mit einfacher Mehrheit. Ein Beherrschungs- oder Gewinnabführungsvertrag ist nicht abgeschlossen. Die Gesellschaft hat eigene Geschäftsführung, Buchhaltung und Personalplanung. Ein Dienstleistungsvertrag mit Leinefaden betrifft die Nutzung eines Servers und einzelne Wartungsleistungen.

3 Havelstrom Speicher AG
Die Havelstrom Speicher AG ist keine Beteiligung der Leinefaden Systeme AG. Zwischen beiden besteht ein Liefervertrag aus dem Jahr 2025. Die Besetzung eines Aufsichtsratsmandats durch Frau Brandt wurde wegen ihrer fachlichen Erfahrung angeregt. Die Aktionäre sind nach den uns mitgeteilten Informationen zwei voneinander unabhängige Familiengesellschaften. Eine gemeinsame Leitung der Unternehmen wurde nicht vereinbart.

4 Grundlage dieser Übersicht
Diese Übersicht wurde aus den internen Beteiligungsunterlagen und den uns überlassenen Angaben erstellt. Sie ist kein Registerauszug. Änderungen nach dem 31. August sind nicht berücksichtigt. Für die Mohnwinkel Service GmbH kann die aktuelle Gesellschafterliste bei deren Geschäftsführung angefordert werden. Eine juristische Konzernbewertung enthält dieses Blatt nicht.

Lennart Volkmann''')
add(A,20,'Variable_Verguetung','eml','Zielvereinbarung 2026 noch offen','Ottilie Rabenau <o.rabenau@leinefaden.example>','2026-09-23T17:15:00+02:00','Dr. Elif Brandt <elif.brandt@postfach.example>', '''Liebe Frau Brandt,

der Aufsichtsrat hat die Zielvereinbarung für dieses Jahr noch nicht beschlossen. Im Juni haben wir nur darüber gesprochen, welche Liquiditäts- und Organisationsziele geeignet wären. Im Juli fehlte Herr Ahlers, im August wurde die Sitzung verschoben. Der Höchstbetrag von 24.000 Euro im Dienstvertrag war als Rahmen gedacht. Eine garantierte Zahlung oder eine monatliche Abschlagszahlung haben wir nach meiner Erinnerung nicht beschlossen.

Ich weiß, dass in der Budgetdatei von Elfriede zwölfmal 2.000 Euro stehen. Das dient dort zur Reservierung der maximalen Jahreskosten und sollte keine Zusage an Sie sein. Bitte reichen Sie der Krankenkasse neben der Datei auch diese Erläuterung ein, falls Sie diese Frage beantworten möchten. Einen Beschluss kann ich Ihnen derzeit nicht beifügen.

In der Oktobersitzung müssen wir klären, wie wir mit der ausgebliebenen Vereinbarung umgehen. Sie haben bereits erhebliche Leistungen erbracht. Ich möchte hierzu vorab aber weder einen Anspruch ablehnen noch einen bestimmten Betrag zusagen. Bitte schicken Sie uns Ihre Darstellung der erreichten Ziele bis zum 6. Oktober.

Mit freundlichen Grüßen
Ottilie Rabenau''')
add(A,21,'Mohnwinkel_Projektstand','docx','Projektstand Lieferantenprozess im September','Dr. Elif Brandt','18. September 2026','Friedrich Knoop, Mohnwinkel Service GmbH', '''1 Bearbeitungsstand
Die Bestandsaufnahme und die Darstellung der Freigabestufen liegen in einer zusammenhängenden Arbeitsfassung vor. Von 25 angeforderten Lieferantenakten wurden 22 vollständig übergeben. In drei Fällen fehlen die Dokumentation der Bonitätsprüfung oder die unterschriebene Freigabe. Die fehlenden Unterlagen sind in einer gesonderten Liste an Herrn Krause benannt. Ohne diese Nachweise wird der Bericht den tatsächlichen Dokumentationsstand beschreiben.

2 Laufende Besprechungen
Die Teilnahme an den Dienstagsrunden hat den Zugang zu Informationen erleichtert. Seit August werden dort jedoch zunehmend Entscheidungen über aktuelle Liefertermine und die Zuordnung von Sachbearbeitern erwartet. Ich habe mehrfach erklärt, dass die endgültige Lieferantenauswahl bei der Geschäftsführung und dem Einkauf bleibt. Bei zwei dringenden Vorgängen habe ich eine Reihenfolge der Rückfragen vorgeschlagen, damit die Arbeiten weitergehen konnten.

3 Verbleibender Umfang
Nach der zuletzt übersandten Zeitaufstellung sind bis Ende August 14,5 Beratungstage verbraucht. Am 8. und 15. September wurden jeweils zwei Stunden für die Abstimmung der Berichtsfassung eingesetzt. Damit verbleibt ein Beratungstag aus dem vereinbarten Kontingent. Diesen plane ich für Endfassung und Abschlussgespräch ein. Weitere operative Besprechungen sind damit nicht eingeplant.

4 Übergabe
Die fertige Fassung wird bis zum 30. September als PDF und bearbeitbare Datei übergeben. Die Auftraggeberin entscheidet anschließend selbst über Umsetzung, Verantwortliche und Fristen. Zugang und Leihgerät gebe ich nach dem Abschlussgespräch zurück. Bitte teilen Sie mir mit, wer die offenen Unterlagen bis zum 24. September bereitstellt.

Dr. Elif Brandt''')
add(A,22,'Chat_Drei_Rollen','txt','Nachrichtenexport zur Septemberabrechnung','Dr. Elif Brandt und Elfriede Holle','21. bis 24. September 2026','Export durch Elfriede Holle', '''21.09.2026 08:18 Elfriede Holle: Die Kasse fragt nach deinem Vertrag. Ist die Aufsichtsratszahlung auch Gehalt bei uns?
21.09.2026 08:24 Elif Brandt: Nein. Havelstrom zahlt selbst am Jahresende. Ich bin dort nicht im Vorstand.
21.09.2026 08:27 Elfriede Holle: In der Exceldatei stand nur Organvergütung. Dann trenne ich die Zeilen.
21.09.2026 08:36 Elif Brandt: Bitte auch Mohnwinkel trennen. Das sind Rechnungen mit Umsatzsteuer, keine Zahlung von Leinefaden.
21.09.2026 08:39 Elfriede Holle: Der Dienstleister meinte, Vorstand sei immer frei. Ich habe nicht nach einzelnen Zweigen gefragt.
21.09.2026 08:44 Elif Brandt: Die Kasse sieht das offenbar genauer. Ich lasse es prüfen, bitte keine Korrektur ohne Rückmeldung.
23.09.2026 15:12 Elfriede Holle: Bonus steht mit 2.000 monatlich im Budget. Das war meine Jahreskostenreserve, bisher nichts ausgezahlt.
23.09.2026 15:14 Elif Brandt: Frau Rabenau schreibt dazu heute noch. Es gibt weiterhin keine unterschriebene Zielvereinbarung.
24.09.2026 09:03 Elfriede Holle: Die Beratungstage aus August waren vier, nicht fünf. Ich hatte einen Havelstromtermin mit eingetragen. Die korrigierte Liste liegt jetzt separat.
24.09.2026 09:08 Elif Brandt: Danke. Bitte den alten Stand nicht löschen, die Kasse hat vermutlich noch den ersten Ausdruck.''')
add(A,23,'Unfallversicherung_Rueckfrage','eml','Versicherungsschutz in den verschiedenen Unternehmen','Zuständige Berufsgenossenschaft <mitgliedschaft@bg-service.example>','2026-09-24T12:10:00+02:00','Dr. Elif Brandt <elif.brandt@postfach.example>', '''Sehr geehrte Frau Brandt,

für die Zuordnung Ihrer Anfrage benötigen wir eine genaue Bezeichnung der jeweiligen Tätigkeit und des Unternehmens. Bitte reichen Sie den Bestellungsbeschluss, den Dienstvertrag und eine kurze Darstellung Ihrer Befugnisse bei Leinefaden ein. Für die Beratungsleistung bei Mohnwinkel benötigen wir den gesonderten Auftrag und Angaben zur tatsächlichen Durchführung. Das Aufsichtsratsmandat ist ebenfalls getrennt zu beschreiben.

Eine Aussage, dass Sie in einem anderen Zweig der Sozialversicherung keine Beiträge zahlen, beantwortet unsere Frage nicht. Für einzelne Tätigkeiten kann zudem eine freiwillige Versicherung in Betracht kommen. Ob hierfür bereits ein Antrag gestellt wurde, können wir anhand Ihrer Nachricht nicht feststellen. In den uns bislang vorliegenden Unterlagen haben wir keine Antragsbestätigung unter Ihrem Namen gefunden.

Bitte teilen Sie außerdem mit, ob Sie eigene Beschäftigte einsetzen und ob ein Unfallereignis Anlass der Anfrage ist. Ihre Nachricht nennt lediglich eine künftige Dienstreise. Eine Entscheidung über einen konkreten Versicherungsfall ist daher derzeit nicht Gegenstand der Bearbeitung.

Mit freundlichen Grüßen
Frieda Niemann, Mitgliedschaft und Beitrag''')
add(A,24,'Erklaerung_weitere_Auftraege','docx','Angaben zu weiteren Tätigkeiten und Einnahmen','Dr. Elif Brandt','27. September 2026','Kanzlei Albers und Semmel', '''1 Selbständige Beratung
Der Auftrag von Mohnwinkel ist mein erster entgeltlicher Beratungsauftrag außerhalb meiner Anstellungen und Organmandate. Ich habe im März 2026 mit den Vorbereitungen begonnen. Ich beschäftige dafür keine Arbeitnehmer und habe auch keine andere Person mit der Bearbeitung beauftragt. Zwei weitere Unternehmen fragten im August unverbindlich an; Verträge wurden daraus nicht. Derzeit rechne ich damit, den Auftrag am 30. September zu beenden.

2 Arbeitsmittel und Durchführung
Die Berichtstexte habe ich auf meinem eigenen Rechner verfasst. Die Einsicht in Lieferantenstammdaten erfolgte auf einem gesicherten Leihgerät der Auftraggeberin. Dienstags war ich häufig vor Ort, aber nicht jede Woche. Die Termine am 12. Mai und 9. Juni habe ich wegen Leinefaden abgesagt und selbst Ersatztermine vorgeschlagen. Einen Urlaubsantrag bei Mohnwinkel habe ich nicht gestellt. Mir wurde jedoch im August gesagt, dass man mich in der Dienstagsrunde fest eingeplant habe.

3 Organmandate
Bei Leinefaden bin ich Vorstandsmitglied. Bei Havelstrom gehöre ich dem Aufsichtsrat an; dort habe ich keine eigene operative Beauftragung erhalten. Ich habe in einer Sitzung Fragen zu Lieferanten gestellt und Unterlagen nachgefordert. Die Anfrage, selbst die Suche zu übernehmen, wurde nicht als Auftrag beschlossen. Meine Vergütung dort wird erst nach Jahresende fällig.

4 Versicherung und Familie
Ich bin seit 2022 freiwilliges Mitglied der Nordhafen Krankenkasse. Meine Kinder wurden am 4. Februar 2008 und am 19. November 2012 geboren. Mein Ehepartner ist selbst gesetzlich krankenversichert. Ich bin in keinem berufsständischen Versorgungswerk und habe keinen Befreiungsbescheid der gesetzlichen Rentenversicherung. Ein freiwilliger Unfallversicherungsantrag liegt mir nicht vor.

Hannover, 27. September 2026
Dr. Elif Brandt''')
add(A,25,'Havelstrom_Verguetungsmitteilung','pdf','Mitteilung zum Aufsichtsratsvergütungskonto','Havelstrom Speicher AG','25. September 2026','Dr. Elif Brandt', '''Sehr geehrte Frau Brandt,

nach dem Hauptversammlungsbeschluss vom 26. März 2026 beträgt Ihre feste Jahresvergütung für das Aufsichtsratsmandat 12.000 Euro. Für den Zeitraum April bis Dezember sind 9.000 Euro vorgemerkt. Die Vergütung ist nach Ablauf des Geschäftsjahres fällig. Bis zum heutigen Tag wurde hierauf keine Zahlung geleistet.

Reisekosten für die Sitzung vom 15. Juni wurden in Höhe von 68,40 Euro am 24. Juni erstattet. Weitere Auslagen wurden uns bislang nicht vorgelegt. Die Erstattung wird auf dem Vergütungskonto getrennt geführt. Ein gesondertes Beratungshonorar ist weder vereinbart noch abgerechnet. Bitte melden Sie uns, falls Ihre eigenen Unterlagen hiervon abweichen.

Der Begriff Organvergütung in unserem Buchhaltungssystem dient der internen Kontierung. Er enthält keine gemeinsame Beurteilung Ihres Vorstandsmandats bei Leinefaden und Ihres Beratungsauftrags bei Mohnwinkel. Wir rechnen ausschließlich die hier genannte Havelstrom-Vergütung ab. Für eine Bescheinigung zu den anderen Gesellschaften wenden Sie sich bitte an deren zuständige Stellen.

Mit freundlichen Grüßen
Waltraud König, Vorsitzende des Aufsichtsrats
Kopie: Buchhaltung''')
add(A,27,'Mohnwinkel_letzte_Mail','eml','Abschluss am Mittwoch','Friedrich Knoop <f.knoop@mohnwinkel.example>','2026-09-28T16:52:00+02:00','Dr. Elif Brandt <elif.brandt@postfach.example>', '''Sehr geehrte Frau Brandt,

wir erwarten Ihre Endfassung am Mittwoch. Die drei fehlenden Aktenstücke konnte unsere Buchhaltung nur teilweise finden. Bitte stellen Sie im Bericht dar, welche Dokumentation tatsächlich vorhanden war. Wir möchten keine nachträglich erzeugten Freigaben als alte Nachweise in die Akte legen. Die Umsetzung übernehmen Herr Krause und sein Team ab Oktober.

Herr Krause hatte Sie bereits für die nächsten Dienstagsrunden eingeladen. Diese Serie wird beendet. Wenn wir später weitere Unterstützung benötigen, sprechen wir über einen neuen Auftrag. Für Oktober ist bislang keine Verlängerung vereinbart. Das Gastkonto bleibt nur bis zum Abschlussgespräch aktiv, danach geben Sie bitte das Leihgerät bei Frau Mertens ab.

Eine Frage der Buchhaltung möchte ich noch weiterreichen: Im Juli steht in Ihrer Tagesliste ein Havelstromtermin, der in der ersten Rechnungsvorbereitung als Mohnwinkelzeit erschien. Nach Ihrer korrigierten Liste wurde er nicht berechnet. Bitte bestätigen Sie das kurz, damit beide Fassungen verständlich zusammen abgelegt werden können. Wir benötigen keine Offenlegung der Inhalte Ihrer anderen Mandate.

Mit freundlichen Grüßen
Friedrich Knoop''')
add(A,28,'Kalender_Erlaeuterung','docx','Erläuterung zur korrigierten Terminübersicht','Dr. Elif Brandt','29. September 2026','Mohnwinkel Service GmbH und eigene Unterlagen', '''1 Korrektur
In meiner ersten Terminübersicht war die Aufsichtsratssitzung der Havelstrom Speicher AG vom 15. Juni versehentlich der Farbe des Mohnwinkelprojekts zugeordnet. Die Zeit wurde in der Rechnungsaufstellung vor Versand der Junirechnung entfernt. Die Rechnung vom 30. Juni über 3,5 Tage enthält die Sitzung nicht. Der Hinweis in Ihrer Nachricht auf einen Julitermin beruht vermutlich auf dem Datum des ersten Excelexports vom 2. Juli.

2 Abgrenzung der Aufzeichnungen
Die Datei enthält Kalendertermine aller drei Tätigkeiten, weil ich meine zeitliche Verfügbarkeit insgesamt planen muss. Ein Kalendereintrag bedeutet nicht automatisch abrechenbare Beratungszeit. In der Rechnungsliste stehen nur die von mir für Mohnwinkel erbrachten und nach dem Vertrag berechneten Leistungen. Die Vorstandsvergütung wird monatlich von Leinefaden abgerechnet. Die Aufsichtsratsvergütung wird hiervon unabhängig von Havelstrom geführt.

3 Umfang der Beratung
Bis Ende August wurden 14,5 Beratungstage erfasst. Im September kamen insgesamt 1,5 Tage für Abstimmung, Endfassung und Übergabe hinzu. Das vereinbarte Kontingent von 16 Tagen ist damit verbraucht. Nicht erfasste kurze Terminabsprachen wurden nicht nachträglich pauschal aufgeschlagen. Einen weiteren Auftrag über laufende Einkaufsleitung habe ich nicht angenommen.

4 Ablage
Bitte bewahren Sie die erste und die korrigierte Übersicht gemeinsam auf. Der erste Export lässt sich dadurch anhand dieser Erklärung nachvollziehen. Die tatsächlichen Rechnungsbeträge werden durch die Korrektur nicht geändert. Die abschließende Septemberrechnung wird nach Übergabe des Berichts erstellt.

Dr. Elif Brandt''')

def docx_record(r,path):
    d=Document(); sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
    for style in d.styles:
        if style.type not in (1,2):continue
        style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        rp=style.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
        for key in list(rf.attrib):
            if 'Theme' in key:del rf.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
        for border in style.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    defaults=d.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        for border in defaults.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(7)
    d.styles['Normal'].paragraph_format.line_spacing=1.12
    if r['case']==S and r['num']==13:
        d.styles['Normal'].paragraph_format.space_after=Pt(5)
        d.styles['Normal'].paragraph_format.line_spacing=1.04
    d.styles['Title'].font.size=Pt(16);d.styles['Title'].font.bold=True
    d.styles['Heading 1'].font.size=Pt(12);d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    d.styles['Heading 1'].paragraph_format.space_after=Pt(7)
    sec.header.paragraphs[0].text=r['author'];sec.header.paragraphs[0].style='Normal'
    d.add_paragraph(r['title'],'Title')
    d.add_paragraph(r['date']+'\n'+r['recipient'])
    import re
    for block in r['body'].split('\n\n'):
        if re.match(r'^\d+(?:\.\d+)* ',block) and '\n' in block:
            h,b=block.split('\n',1);d.add_paragraph(h,'Heading 1');d.add_paragraph(b)
        else:d.add_paragraph(block)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    footer.add_run('Seite ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    d.core_properties.author=r['author'].split('<')[0].strip();d.core_properties.title=r['title']
    d.core_properties.subject='';d.core_properties.keywords=''
    for p in d.paragraphs:
        for border in p._p.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    d.save(path)

def pdf_record(r,path):
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=60,leftMargin=60,topMargin=54,bottomMargin=54,title=r['title'],author=r['author'])
    sty=ParagraphStyle('body',fontName='Times-Roman',fontSize=11,leading=14.5,spaceAfter=9)
    head=ParagraphStyle('title',parent=sty,fontName='Times-Bold',fontSize=15,leading=18,spaceAfter=16)
    story=[Paragraph(escape(r['author']),sty),Spacer(1,10),Paragraph(escape(r['title']),head),Paragraph(escape(r['date'])+'<br/>'+escape(r['recipient']),sty),Spacer(1,9)]
    for block in r['body'].split('\n\n'):story.append(Paragraph(escape(block).replace('\n','<br/>'),sty))
    def footer(c,d):c.setFont('Times-Roman',9);c.drawRightString(A4[0]-60,32,f'Seite {d.page}')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

def eml_record(r,path):
    m=EmailMessage(policy=SMTP);m['From']=r['author'];m['To']=r['recipient'];m['Subject']=r['title'];m['Date']=datetime.fromisoformat(r['date']);m['Message-ID']=f"<{r['case']}.{r['num']:02d}@aktenpost.example>"
    m.set_content(r['body'],charset='utf-8');path.write_bytes(m.as_bytes())

def attach_evidence():
    from email import policy
    from email.parser import BytesParser
    mappings={(S,1):[3,8,13,16],(S,10):[8,29],(S,14):[7,11],(S,27):[26]}
    for (case,num),numbers in mappings.items():
        folder=ROOT/'testakten'/case
        mail=next(folder.glob(f'{num:02d}_*.eml'))
        msg=BytesParser(policy=policy.SMTP).parsebytes(mail.read_bytes())
        for n in numbers:
            source=next(p for p in folder.glob(f'{n:02d}_*') if p.suffix in ['.pdf','.docx','.csv'])
            subtype={'.pdf':'pdf','.docx':'vnd.openxmlformats-officedocument.wordprocessingml.document','.csv':'csv'}[source.suffix]
            msg.add_attachment(source.read_bytes(),maintype='text' if source.suffix=='.csv' else 'application',subtype=subtype,filename=source.name)
        mail.write_bytes(msg.as_bytes())

def run():
    for c in [S,A]:(ROOT/'testakten'/c).mkdir(parents=True,exist_ok=True)
    for r in records:
        path=ROOT/'testakten'/r['case']/f"{r['num']:02d}_{r['stem']}.{r['kind']}"
        if r['kind']=='docx':docx_record(r,path)
        elif r['kind']=='pdf':pdf_record(r,path)
        elif r['kind']=='eml':eml_record(r,path)
        else:path.write_text(r['title']+'\n'+r['date']+'\n'+r['author']+'\n\n'+r['body']+'\n',encoding='utf-8')
    with (ROOT/'testakten'/S/'26_Kanzlei_Zahlungseingaenge.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f,delimiter=';');w.writerow(['Zahlungstag','Rechnung','Mandat','Netto_EUR','Umsatzsteuer_EUR','Brutto_EUR','Notiz'])
        w.writerows([['2026-02-11','MA-26-01','Nachbarschaft 01',450,85.5,535.5,'Beratung'],['2026-04-17','MA-26-03','Mietvertrag 02',680,129.2,809.2,'Vertragsprüfung'],['2026-06-22','MA-26-04','Kaufvertrag 03',920,174.8,1094.8,'Entwurf'],['2026-09-03','MA-26-05','Nachbarschaft 04',1200,228,1428,'Rechnung vom 28. Juli; Zahlung erst September']])
    with (ROOT/'testakten'/A/'26_Termine_korrigiert.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f,delimiter=';');w.writerow(['Datum','Gesellschaft','Termin','Stunden','Abgerechnet_im_Beratungsprojekt','Bemerkung'])
        w.writerows([['2026-06-02','Mohnwinkel','Interview Einkauf',4,'ja','Tagesliste'],['2026-06-09','Leinefaden','Vorstandsbericht',5,'nein','Mohnwinkeltermin abgesagt'],['2026-06-15','Havelstrom','Aufsichtsrat',4,'nein','Farbzuordnung im ersten Export falsch'],['2026-06-16','Mohnwinkel','Aktenstichprobe',8,'ja','Vor Ort'],['2026-08-11','Mohnwinkel','Freigabeprozess',8,'ja','Leihgerät genutzt'],['2026-09-08','Mohnwinkel','Berichtsabstimmung',2,'ja','Videokonferenz'],['2026-09-15','Mohnwinkel','Berichtsabstimmung',2,'ja','Videokonferenz']])
    warn='Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\nThis test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'
    for c,title in [(S,'Mila Ahrens und Fleetbogen'),(A,'Dr. Elif Brandt und drei Gesellschaften')]:
        text=f'# {title}\n\nArbeitsakte mit Vertragsunterlagen, Korrespondenz, Bescheiden und Abrechnungsdaten. Stand der zusammengestellten Unterlagen ist der 30. September 2026.\n\n<!-- reserved-example-contacts -->\nKontaktadressen mit der reservierten Endung `.example` sind erfunden und dienen ausschließlich der sicheren Darstellung der Akte.\n\n{warn}\n\n| Fassung | Download |\n| --- | --- |\n| Gesamt-PDF | [Gesamtakte lesen](gesamt-pdf/{c}_gesamt.pdf) |\n| Originaldateien | [Akten-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{c}.zip) |\n| Einzelne PDFs | [Einzel-PDF-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{c}-einzelpdfs.zip) |\n'
        (ROOT/'testakten'/c/'README.md').write_text(text,encoding='utf-8')
    attach_evidence()
    (TMP/'record-data.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    manifest=[]
    for c in [S,A]:
        for p in sorted((ROOT/'testakten'/c).glob('*')):
            if p.is_file():manifest.append({'case':c,'name':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    (TMP/'native-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'records':len(records),'cases':[S,A],'manifest':str(TMP/'native-manifest.json')},ensure_ascii=False))

if __name__=='__main__':run()
