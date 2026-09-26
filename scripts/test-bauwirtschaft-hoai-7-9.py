#!/usr/bin/env python3
"""Offline-Regression der Originale; optional kanonische Exporte. Autor: Klotzkette."""
from __future__ import annotations

import argparse
import csv
from email import policy
from email.parser import BytesParser
import importlib.util
import hashlib
import io
from pathlib import Path
import re
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageStat
from pypdf import PdfReader
import yaml
from readme_decimal_headings import normalize_decimal_headings
from testakte_disclaimer import NOTICE_BYTES, NOTICE_DE, NOTICE_EN, NOTICE_MARKDOWN

HERE=Path(__file__).resolve().parent
spec=importlib.util.spec_from_file_location("hoai789_build",HERE/"build-bauwirtschaft-hoai-7-9-akten.py")
B=importlib.util.module_from_spec(spec);sys.modules[spec.name]=B;spec.loader.exec_module(B);B.define_all()
parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("--assets",type=Path)
args,unknown=parser.parse_known_args();ASSETS=args.assets
SLUGS={7:"hoai-7-angebote-werten-und-vergabe-vorbereiten",8:"hoai-8-bauueberwachung-und-dokumentation-fuehren",9:"hoai-9-objektbetreuung-und-maengelverfolgung-organisieren"}
XLSX={7:(6,"Planer-LV",{"F5":39736,"F7":47285.84}),8:(12,"Kostenjournal",{"D5":394586,"E6":377840,"G7":448964.6}),9:(12,"Sicherungen",{"D10":4641,"D11":2844.1,"F5":7485.1})}


def pdf_text(path):return " ".join(" ".join(p.extract_text() or "" for p in PdfReader(path).pages).split())


class Originals(unittest.TestCase):
    def test_exact_inventory_and_ascii_names(self):
        for phase,count in [(7,16),(8,20),(9,16)]:
            folder=B.ROOT/"testakten"/B.CASES[phase]["slug"]
            expected={d["filename"] for d in B.FILES[phase].values()}
            actual={p.name for p in folder.iterdir() if p.is_file() and p.name not in {"README.md","rubric.yaml"}}
            self.assertEqual(actual,expected);self.assertEqual(len(expected),count)
            for name in expected:
                self.assertTrue(name.isascii(),name);self.assertGreater((folder/name).stat().st_size,100,name)
            self.assertTrue({".docx",".pdf",".xlsx",".png",".eml",".csv"}.issubset({Path(n).suffix for n in expected}))

    def test_frontmatter_sources_and_standalone_prompts(self):
        for phase,slug in SLUGS.items():
            path=B.ROOT/"bauwirtschaft/skills"/slug/"SKILL.md";text=path.read_text();meta=yaml.safe_load(text.split("---",2)[1])
            self.assertEqual(set(meta),{"name","description"});self.assertEqual(meta["name"],slug)
            self.assertTrue(80<=len(meta["description"])<=1024);self.assertLess(len(text.splitlines()),500)
            self.assertEqual(len(re.findall(r"^## [1-6]\. ",text,re.M)),6)
            self.assertNotIn("werkstatt.md",text)
            for link in re.findall(r"\]\(([^)]+)\)",text):
                if not link.startswith("https://"):self.assertTrue((path.parent/link).resolve().is_file(),link)
            prompt=(B.ROOT/f"bauwirtschaft/bauwirtschaft-hoai-{phase}-werkstatt.md").read_text()
            self.assertEqual(len(re.findall(r"^# ",prompt,re.M)),1)
            self.assertGreater(len(prompt.encode()),20000)
            self.assertNotRegex(prompt,r"\]\((?:\.\./|references/)")
            self.assertNotIn(chr(167),text+prompt)

    def test_readmes_and_rubric(self):
        for phase,c in B.CASES.items():
            folder=B.ROOT/"testakten"/c["slug"];text=(folder/"README.md").read_text()
            self.assertEqual(text,normalize_decimal_headings(text));self.assertIn(f"Leistungsphase {phase}",text.split("# ",1)[1].splitlines()[0])
            self.assertIn(NOTICE_MARKDOWN+"\n\n| Was",text);self.assertIn("erfunden",text)
            self.assertIn("<!-- reserved-example-contacts -->",text)
            self.assertIn("reservierte `.example`-Domains",text)
            self.assertEqual(text.count("<!-- BEGIN gesamt-pdf-section (autogen) -->"),1)
            self.assertEqual(text.count("<!-- END gesamt-pdf-section (autogen) -->"),1)
            self.assertEqual(text.count("| Was | Format | Quelle |"),1)
            spec=importlib.util.spec_from_file_location("hoai789_inject",HERE/"inject-gesamt-pdf-section.py")
            injector=importlib.util.module_from_spec(spec);spec.loader.exec_module(injector)
            with tempfile.TemporaryDirectory(prefix="hoai789-inject-") as tmp:
                temp=Path(tmp);readme=temp/"README.md";readme.write_text(text,encoding="utf-8")
                (temp/"01_Original.pdf").touch()
                if (folder/"gesamt-pdf"/(c["slug"]+"_gesamt.pdf")).exists():
                    (temp/"gesamt-pdf").mkdir();(temp/"gesamt-pdf"/(c["slug"]+"_gesamt.pdf")).touch()
                self.assertEqual(injector.inject(readme,c["slug"]),"unchanged")
            rubric=yaml.safe_load((folder/"rubric.yaml").read_text());self.assertEqual(rubric["plugin"],"bauwirtschaft")
            ids=[c["id"] for c in rubric["checks"]];self.assertEqual(len(ids),len(set(ids)))
            self.assertTrue(any(c["check_type"]=="human_review" for c in rubric["checks"]))

    def test_document_text_and_metadata(self):
        for phase in B.CASES:
            for d in B.FILES[phase].values():
                path=B.doc_path(d);ext=path.suffix;text=""
                if ext==".pdf":
                    reader=PdfReader(path);text=pdf_text(path);self.assertEqual(reader.metadata.author,B.AUTHOR)
                    self.assertGreater(len(text),600,path.name)
                elif ext==".docx":
                    doc=Document(path);self.assertEqual(doc.core_properties.author,B.AUTHOR)
                    text="\n".join(p.text for p in doc.paragraphs);self.assertGreater(len(text),600,path.name)
                    self.assertEqual(doc.styles["Normal"].font.name,"Times New Roman")
                elif ext==".txt":text=path.read_text();self.assertGreater(len(text),600,path.name)
                if text:
                    self.assertNotIn(NOTICE_DE,text);self.assertNotIn(NOTICE_EN,text);self.assertNotIn(chr(167),text)
                    self.assertIn("gez.",text,path.name)

    def test_mime_attachments_are_actual_originals(self):
        count=0
        for phase in B.CASES:
            for d in B.FILES[phase].values():
                path=B.doc_path(d)
                if path.suffix!=".eml":continue
                msg=BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                for field in ["From","To","Date","Subject","Message-ID"]:self.assertTrue(msg[field],(path.name,field))
                self.assertIn(".example",msg["From"]);body=msg.get_body(preferencelist=("plain",)).get_content()
                self.assertGreater(len(body),450);self.assertNotIn(NOTICE_DE,body)
                parts=list(msg.iter_attachments());self.assertEqual(len(parts),len(d["attachments"]))
                for part in parts:
                    self.assertEqual(part.get_payload(decode=True),(path.parent/part.get_filename()).read_bytes());count+=1
        self.assertEqual(count,3)

    def test_cached_recalculated_workbooks(self):
        for phase,(number,sheet,expected) in XLSX.items():
            path=B.doc_path(B.FILES[phase][number]);formulas=load_workbook(path,data_only=False);values=load_workbook(path,data_only=True)
            self.assertEqual(formulas.properties.creator,B.AUTHOR);self.assertEqual(formulas.properties.lastModifiedBy,B.AUTHOR)
            for cell,value in expected.items():
                self.assertEqual(formulas[sheet][cell].data_type,"f",(phase,cell))
                self.assertAlmostEqual(values[sheet][cell].value,value,places=2)
                self.assertIn("407",formulas[sheet][cell].number_format)
            for row in values[sheet]:
                for cell in row:self.assertNotEqual(cell.data_type,"e")
            self.assertEqual(formulas[sheet].sheet_view.showGridLines,False)
            formulas.close();values.close()

    def test_raster_plans_nonblank_and_metadata(self):
        for phase in B.CASES:
            for d in B.FILES[phase].values():
                path=B.doc_path(d)
                if path.suffix!=".png":continue
                with Image.open(path) as im:
                    self.assertGreaterEqual(im.width,1500);self.assertGreaterEqual(im.height,1000);self.assertEqual(im.info.get("Author"),B.AUTHOR)
                    std=ImageStat.Stat(im.convert("RGB")).stddev;self.assertGreater(min(std),20)
                    colors=im.resize((200,150)).getcolors(30000);self.assertGreater(len(colors or []),50)

    def test_case_arithmetic_and_evidence(self):
        self.assertEqual(sum(q*p for q,p in zip(B.Q7,B.EP7)),39736)
        self.assertEqual(sum(q*p for q,p in zip(B.Q8,B.EP8)),36986)
        self.assertEqual(18*12,216);self.assertEqual(6*(18+72+39),774)
        text=pdf_text(B.doc_path(B.FILES[8][8]));self.assertIn("39.640,00",text);self.assertIn("29.750,00",text)
        self.assertIn("nicht vorgenommen",pdf_text(B.doc_path(B.FILES[8][6])) if B.doc_path(B.FILES[8][6]).suffix==".pdf" else " ".join(p.text for p in Document(B.doc_path(B.FILES[8][6])).paragraphs))
        self.assertIn("nicht sichtbar",pdf_text(B.doc_path(B.FILES[8][16])))
        self.assertIn("keine erneute Abnahme",pdf_text(B.doc_path(B.FILES[9][3])))
        self.assertIn("78",pdf_text(B.doc_path(B.FILES[9][9])));self.assertIn("ohne Einheit",pdf_text(B.doc_path(B.FILES[9][9])))
        for phase,num in [(7,15),(8,5),(8,11),(9,13)]:
            with B.doc_path(B.FILES[phase][num]).open(encoding="utf-8-sig",newline="") as f:rows=list(csv.reader(f))
            self.assertGreater(len(rows),3);self.assertTrue(all(len(r)==len(rows[0]) for r in rows))

    @unittest.skipUnless(ASSETS,"Kanonische Exporte nur mit --assets DIR")
    def test_canonical_exports(self):
        for phase,c in B.CASES.items():
            slug=c["slug"];folder=B.ROOT/"testakten"/slug
            for suffix,originals in [("",True),("-einzelpdfs",False)]:
                path=ASSETS/f"testakte-{slug}{suffix}.zip";self.assertTrue(path.is_file(),str(path))
                with zipfile.ZipFile(path) as z:
                    names=z.namelist();self.assertEqual(len(names),len(set(names)));self.assertTrue(all("/" not in n for n in names))
                    self.assertFalse(any(n.endswith((".md",".yaml")) for n in names));notice=z.read("README.txt").decode("utf-8")
                    self.assertEqual(z.read("README.txt"),NOTICE_BYTES)
                    self.assertIn(NOTICE_DE,notice[:500]);self.assertIn(NOTICE_EN,notice[:500])
                    for d in B.FILES[phase].values():
                        name=d["filename"] if originals else Path(d["filename"]).stem+".pdf"
                        self.assertIn(name,names)
                        if originals:self.assertEqual(hashlib.sha256(z.read(name)).hexdigest(),hashlib.sha256((folder/name).read_bytes()).hexdigest(),f"{slug}/{name}: Exportstand weicht vom Original ab")
                        else:
                            text=pdf_text(io.BytesIO(z.read(name)));self.assertNotIn(NOTICE_DE,text);self.assertNotIn(NOTICE_EN,text)
                    if not originals:self.assertEqual(len(names),len(B.FILES[phase])+1)
            total=folder/"gesamt-pdf"/(slug+"_gesamt.pdf");self.assertTrue(total.is_file())
            text=pdf_text(total);self.assertNotIn(NOTICE_DE,text);self.assertNotIn(NOTICE_EN,text)


if __name__=="__main__":unittest.main(argv=[sys.argv[0],*unknown])
