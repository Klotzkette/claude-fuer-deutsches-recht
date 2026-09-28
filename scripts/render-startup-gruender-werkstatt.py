#!/usr/bin/env python3
"""Erstellt DOCX und PDF als Lesefassungen des kanonischen Werkstatt-Prompts.

A4, Times New Roman 11 pt, natürliche Seitenumbrüche, echte Quellenlinks und
PDF-Lesezeichen. Die Seitenzahl wird gemessen und nicht durch Leerumbrüche
vorgegeben. PNG-Seiten für die anschließende Sichtprüfung liegen getrennt.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import zipfile
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, ByteStringObject, NameObject

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'startup-gruender/startup-gruender-werkstatt.md'
OUT = ROOT / 'startup-gruender/assets'
LINK = re.compile(r'\[([^\]]+)\]\((https?://[^\s)]+)\)')
DATE = datetime(2026, 9, 28, 9, 0, tzinfo=timezone.utc)


def inline(paragraph, text):
    cursor = 0
    for match in LINK.finditer(text):
        paragraph.add_run(text[cursor:match.start()])
        relationship = paragraph.part.relate_to(match.group(2), RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), relationship)
        hyperlink.set(qn('w:history'), '1')
        run = OxmlElement('w:r')
        properties = OxmlElement('w:rPr')
        fonts = OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'), 'Times New Roman')
        fonts.set(qn('w:hAnsi'), 'Times New Roman')
        properties.append(fonts)
        size = OxmlElement('w:sz'); size.set(qn('w:val'), '22'); properties.append(size)
        color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79'); properties.append(color)
        underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); properties.append(underline)
        run.append(properties)
        node = OxmlElement('w:t'); node.text = match.group(1); run.append(node)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def build_docx(text: str, path: Path):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.4)
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.header_distance = s.footer_distance = Cm(1.2)
    normal = d.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 20), ('Heading 1', 14)]:
        style = d.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name == 'Heading 1'
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_before = Pt(16 if name == 'Heading 1' else 0)
        style.paragraph_format.space_after = Pt(9)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    for node in list(d.styles.element.iter()):
        if node.tag == qn('w:pBdr'):
            node.getparent().remove(node)
        elif node.tag == qn('w:rFonts'):
            for key in list(node.attrib):
                if 'theme' in key.lower(): del node.attrib[key]
            node.set(qn('w:ascii'), 'Times New Roman')
            node.set(qn('w:hAnsi'), 'Times New Roman')
    d.core_properties.author = 'Startup-Gründer'
    d.core_properties.title = 'Startup Gründer Werkstatt'
    d.core_properties.subject = 'Lesefassung des Werkstatt-Prompts'
    d.core_properties.created = d.core_properties.modified = DATE
    heading_count = 0
    for block in re.split(r'\n\s*\n', text.strip()):
        if block.startswith('# '):
            d.add_paragraph(block[2:].strip(), 'Title')
        elif block.startswith('## '):
            heading_count += 1
            p = d.add_paragraph(block[3:].strip(), 'Heading 1')
            bookmark = OxmlElement('w:bookmarkStart')
            bookmark.set(qn('w:id'), str(heading_count))
            bookmark.set(qn('w:name'), f'Kapitel_{heading_count}')
            p._p.insert(0, bookmark)
            end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(heading_count)); p._p.append(end)
        else:
            if block.startswith(('#', '|', '```', '- ', '* ')):
                raise ValueError('Nicht unterstütztes Markdown-Element; Renderer gezielt erweitern.')
            p = d.add_paragraph()
            inline(p, ' '.join(block.splitlines()))
    if heading_count != 40:
        raise ValueError(f'Erwartet 40 Workflow-Kapitel, gefunden {heading_count}.')
    footer = s.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run('Startup Gründer Werkstatt · ')
    run.font.name = 'Times New Roman'; run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 90, 90)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    stream = io.BytesIO(); d.save(stream)
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 28, 9, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))


def normalize_pdf(source: Path, target: Path, source_hash: str):
    reader = PdfReader(source)
    writer = PdfWriter(); writer.clone_document_from_reader(reader)
    writer.add_metadata({'/Title':'Startup Gründer Werkstatt','/Author':'Startup-Gründer','/Subject':'Lesefassung des Werkstatt-Prompts','/CreationDate':'D:20260928090000Z','/ModDate':'D:20260928090000Z'})
    if '/Metadata' in writer._root_object:
        del writer._root_object[NameObject('/Metadata')]
    identifier = bytes.fromhex(source_hash[:32])
    writer._ID = ArrayObject([ByteStringObject(identifier), ByteStringObject(identifier)])
    with target.open('wb') as stream: writer.write(stream)


def renderer_environment(qa_dir: Path, runtime: Path):
    """Expose locally installed Times New Roman to bundled Fontconfig on macOS."""
    env = os.environ.copy()
    font_dir = Path('/System/Library/Fonts/Supplemental')
    if sys.platform == 'darwin' and (font_dir / 'Times New Roman.ttf').is_file():
        bundled_config = runtime / 'dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fontconfig/fonts.conf'
        config = qa_dir.resolve() / 'fonts.conf'
        cache = qa_dir.resolve() / 'font-cache'
        config.write_text('<?xml version="1.0"?>\n<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">\n'
                          '<fontconfig><dir>' + escape(str(font_dir)) + '</dir><cachedir>' + escape(str(cache)) +
                          '</cachedir><include ignore_missing="yes">' + escape(str(bundled_config)) +
                          '</include></fontconfig>\n', encoding='utf-8')
        env['FONTCONFIG_FILE'] = str(config)
        env['SAL_FONTPATH'] = str(font_dir)
    return env


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--source', type=Path, default=SOURCE)
    p.add_argument('--output-dir', type=Path, default=OUT)
    p.add_argument('--qa-dir', type=Path, default=Path('/tmp/startup-workshop-qa'))
    p.add_argument('--renderer', type=Path)
    args = p.parse_args()
    text = args.source.read_text(encoding='utf-8')
    source_hash = hashlib.sha256(args.source.read_bytes()).hexdigest()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.qa_dir.mkdir(parents=True, exist_ok=True)
    target_docx = args.output_dir / 'startup-gruender-werkstatt.docx'
    target_pdf = args.output_dir / 'startup-gruender-werkstatt.pdf'
    build_docx(text, target_docx)
    runtime = Path(sys.executable).resolve().parents[3]
    renderer = args.renderer or runtime / 'plugins/openai-primary-runtime/plugins/documents/skills/documents/render_docx.py'
    if not renderer.is_file():
        raise SystemExit('Gebündelten render_docx.py über --renderer angeben; kein Desktop-LibreOffice-Fallback.')
    render_dir = args.qa_dir / 'render'
    render_dir.mkdir(exist_ok=True)
    for stale_page in render_dir.glob('page-*.png'): stale_page.unlink()
    result = subprocess.run([sys.executable,str(renderer),str(target_docx),'--output_dir',str(render_dir),'--emit_pdf','--dpi','120'],env=renderer_environment(args.qa_dir, runtime),text=True,capture_output=True,timeout=300)
    (args.qa_dir/'render.log').write_text(result.stdout+result.stderr,encoding='utf-8')
    if result.returncode:
        raise SystemExit('Rendering fehlgeschlagen; Einzelheiten stehen in render.log.')
    normalize_pdf(render_dir/(target_docx.stem+'.pdf'),target_pdf,source_hash)
    reader = PdfReader(target_pdf)
    pages = len(reader.pages)
    pdf_fonts = sorted({str(font.get_object().get('/BaseFont')) for page in reader.pages for font in page['/Resources']['/Font'].values()})
    links = LINK.findall(text)
    result = {'source':str(args.source.relative_to(ROOT)) if args.source.is_relative_to(ROOT) else str(args.source), 'source_sha256':source_hash,'source_words':len(text.split()),'workflow_chapters':40,'source_links':len(links),'pages':pages,'manual_page_breaks':0,'font':'Times New Roman','body_points':11,'line_spacing':1.5,'a4_margins_cm':{'left':2.4,'right':2.4,'top':2.2,'bottom':2.2},'docx_sha256':hashlib.sha256(target_docx.read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256(target_pdf.read_bytes()).hexdigest(),'visual_review':'Noch durchzuführen; Renderer bestätigt nur Erstellung.'}
    result['pdf_embedded_fonts'] = pdf_fonts
    result['pdf_uses_times_new_roman'] = bool(pdf_fonts) and all('TimesNewRoman' in name for name in pdf_fonts)
    (args.qa_dir/'render-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
