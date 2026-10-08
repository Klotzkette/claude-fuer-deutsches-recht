"""Lesbare PDF-Darstellung ausschließlich der drei AML-Fallregister.

Die XLSX-Datei bleibt unverändert. Formelergebnisse stammen aus ihren geprüften
Caches; ohne vorhandenen Cache wird kein vermeintlich vollständiges PDF erzeugt.
"""

from __future__ import annotations

from datetime import date, datetime
import io
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle,
)

CASES = frozenset({
    "aml-unternehmen-werkzeughandel-erfurt",
    "aml-kanzlei-grundstueck-berlin",
    "aml-notariat-kaufpreis-wuerzburg",
})
MARGIN = 42
WIDTH = A4[0] - 2 * MARGIN


def matches(path: Path) -> bool:
    parts = path.parts
    return (
        len(parts) >= 4 and parts[-4] == "testakten" and parts[-3] in CASES
        and parts[-2:] == ("04_Tabellen", "Fallregister.xlsx")
    )


def _fonts() -> tuple[str, str, str]:
    choices = [
        (Path("/System/Library/Fonts/Supplemental"), "Times New Roman.ttf", "Times New Roman Bold.ttf", "Times New Roman"),
        (Path("C:/Windows/Fonts"), "times.ttf", "timesbd.ttf", "Times New Roman"),
        (Path("/usr/share/fonts/truetype/msttcorefonts"), "Times_New_Roman.ttf", "Times_New_Roman_Bold.ttf", "Times New Roman"),
        (Path("/usr/share/fonts/truetype/liberation2"), "LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "Liberation Serif"),
        (Path("/usr/share/fonts/truetype/liberation"), "LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "Liberation Serif"),
    ]
    for directory, regular, bold, label in choices:
        if (directory / regular).is_file() and (directory / bold).is_file():
            family = "AMLTimesNewRoman" if label == "Times New Roman" else "AMLLiberationSerif"
            if family not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(family, str(directory / regular)))
                pdfmetrics.registerFont(TTFont(family + "Bold", str(directory / bold)))
            return family, family + "Bold", label
    raise ValueError("AML-PDF benötigt Times New Roman oder Liberation Serif (fonts-liberation2).")


def _text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        return "WAHR" if value else "FALSCH"
    if isinstance(value, (int, float)):
        fmt = cell.number_format
        percent = "%" in fmt
        decimals = 2 if ".00" in fmt else (1 if ".0" in fmt else 0)
        if percent:
            value *= 100
        if decimals or "#,##" in fmt:
            result = f"{value:,.{decimals}f}".replace(",", "_").replace(".", ",").replace("_", ".")
        else:
            result = f"{value:g}"
        return result + (" %" if percent else "")
    return str(value)


def _table(rows, widths, repeat=0):
    table = LongTable(rows, colWidths=widths, repeatRows=repeat, splitByRow=1, splitInRow=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOX", (0, 0), (-1, -1), 0.45, colors.HexColor("#8A949A")),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C4CDD1")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        *( [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF0"))] if repeat else [] ),
    ]))
    return table


def render(path: Path) -> bytes | None:
    """Rendert bekannte AML-Register vollständig; fremde Pfade liefern None."""
    path = Path(path)
    if not matches(path):
        return None
    source = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    try:
        if values.sheetnames != ["Übersicht", "Zahlungen", "Beteiligte", "Verlauf"]:
            raise ValueError(f"{path}: unerwartete Tabellenblätter")
        for sheet in source:
            for row in sheet:
                for cell in row:
                    cached = values[sheet.title][cell.coordinate]
                    if cell.data_type == "f" and cached.value is None:
                        raise ValueError(f"{sheet.title}!{cell.coordinate}: Formelcache fehlt")
                    if cached.data_type == "e":
                        raise ValueError(f"{sheet.title}!{cell.coordinate}: Excel-Fehler {cached.value}")

        regular, bold, font_label = _fonts()
        styles = {
            "body": ParagraphStyle("AMLBody", fontName=regular, fontSize=11, leading=14, spaceAfter=6, splitLongWords=True),
            "cell": ParagraphStyle("AMLCell", fontName=regular, fontSize=9.5, leading=12, splitLongWords=True),
            "bold": ParagraphStyle("AMLCellBold", fontName=bold, fontSize=9.5, leading=12, splitLongWords=True),
            "heading": ParagraphStyle("AMLHeading", fontName=bold, fontSize=14, leading=17, spaceAfter=10, alignment=TA_LEFT),
        }

        def paragraph(text, style="cell"):
            return Paragraph(escape(text).replace("\n", "<br/>"), styles[style])

        flow = []
        for index, sheet in enumerate(values):
            if index:
                flow.append(PageBreak())
            flow.append(paragraph(f"{index + 1}. {sheet.title}", "heading"))
            if not index and font_label != "Times New Roman":
                flow.append(paragraph("Schrift: Liberation Serif als Ersatz für Times New Roman.", "body"))

            if sheet.title == "Zahlungen":
                header = None
                for row in sheet:
                    entries = [_text(cell) for cell in row]
                    occupied = [text for text in entries if text]
                    if not occupied:
                        continue
                    if len(occupied) == 1:
                        flow.append(paragraph(occupied[0], "body"))
                        continue
                    if header is None:
                        if len(entries) != 9 or entries[0] != "Vorgang-ID":
                            raise ValueError("Unerwarteter Aufbau des AML-Zahlungsregisters")
                        header = entries
                        continue
                    if len(entries) != 9:
                        raise ValueError("AML-Zahlungsregister muss genau neun Spalten enthalten")
                    head = _table([
                        [paragraph(text, "bold") for text in header[:4]],
                        [paragraph(text) for text in entries[:4]],
                    ], [WIDTH * share for share in (0.17, 0.20, 0.20, 0.43)], repeat=1)
                    details = _table([
                        [paragraph(label, "bold"), paragraph(text)]
                        for label, text in zip(header[4:], entries[4:])
                    ], [WIDTH * 0.28, WIDTH * 0.72])
                    flow.append(KeepTogether([head, details, Spacer(1, 10)]))
                if header is None:
                    raise ValueError("AML-Zahlungsregister hat keine Kopfzeile")
            else:
                if sheet.max_column != 4:
                    raise ValueError(f"{sheet.title}: erwartet werden vier Spalten")
                widths = {
                    "Übersicht": (0.22, 0.16, 0.12, 0.50),
                    "Beteiligte": (0.20, 0.23, 0.34, 0.23),
                    "Verlauf": (0.15, 0.45, 0.14, 0.26),
                }[sheet.title]
                group = []
                repeat = 0

                def flush():
                    nonlocal group, repeat
                    if group:
                        flow.append(_table(group, [WIDTH * share for share in widths], repeat))
                        flow.append(Spacer(1, 7))
                    group, repeat = [], 0

                for row in sheet:
                    occupied = [cell for cell in row if cell.value is not None and str(cell.value) != ""]
                    if not occupied:
                        flush()
                        continue
                    if len(occupied) == 1:
                        flush()
                        flow.append(paragraph(_text(occupied[0]), "body"))
                        continue
                    if not group:
                        repeat = int(len(occupied) == 4 and all(cell.font.bold for cell in occupied))
                    group.append([paragraph(_text(cell), "bold" if cell.font.bold else "cell") for cell in row])
                flush()

        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=MARGIN, rightMargin=MARGIN,
                                topMargin=58, bottomMargin=45, title=f"Fallregister - {path.parts[-3]}",
                                author="Kanzleiakte", pageCompression=1)

        def furniture(canvas, document):
            canvas.saveState()
            canvas.setFont(regular, 9)
            canvas.drawString(MARGIN, A4[1] - 31, path.parts[-3])
            canvas.drawRightString(A4[0] - MARGIN, 26, f"Seite {document.page}")
            canvas.restoreState()

        def stable_canvas(*args, **kwargs):
            kwargs["invariant"] = 1
            return Canvas(*args, **kwargs)

        doc.build(flow, onFirstPage=furniture, onLaterPages=furniture, canvasmaker=stable_canvas)
        return output.getvalue()
    finally:
        source.close()
        values.close()
