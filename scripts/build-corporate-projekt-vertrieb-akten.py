#!/usr/bin/env python3
"""Erzeugt ausschließlich die nativen Quellen der beiden Projekt- und Vertriebsakten."""

import argparse
import csv
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from akten_docx_format import separate_section_headings
from PIL import Image, ImageDraw
from akten_build_runtime import node_binary, screen_font
from testakte_office_pdf import render_office_batch


ROOT = Path(__file__).resolve().parents[1]
P = ROOT / 'testakten/corporate-contract-law-projektvertrag-automation-augsburg'
V = ROOT / 'testakten/corporate-contract-law-vertrieb-messtechnik-bremen'
NOTICE = ('Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.',
          'This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.')
FILES = {P: [], V: []}


def remember(folder, name):
    FILES[folder].append(name)
    folder.mkdir(parents=True, exist_ok=True)
    return folder / name


def doc(folder, name, author, title, meta, body):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    section.header_distance, section.footer_distance = Cm(0.9), Cm(0.9)
    for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
        style = document.styles[style_name]
        style.font.name = 'Times New Roman'
        run_fonts = style.element.get_or_add_rPr().rFonts
        for attribute in list(run_fonts.attrib):
            if attribute.endswith('Theme'):
                del run_fonts.attrib[attribute]
        for script in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            run_fonts.set(qn(f'w:{script}'), 'Times New Roman')
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(4)
    for node in document.styles.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    document.styles['Normal'].paragraph_format.line_spacing = 1.0
    document.styles['Title'].font.size = Pt(17)
    document.styles['Title'].font.bold = True
    for s in ['Heading 1', 'Heading 2']:
        document.styles[s].font.size = Pt(12 if s == 'Heading 1' else 11)
        document.styles[s].font.bold = True
        document.styles[s].paragraph_format.space_before = Pt(8)
        document.styles[s].paragraph_format.space_after = Pt(0)
    if name == '10_Schnittstellenbeiblatt_SB-03.docx':
        section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.5)
        document.styles['Normal'].paragraph_format.space_after = Pt(2)
        for s in ['Heading 1', 'Heading 2']:
            document.styles[s].paragraph_format.space_before = Pt(6)
            document.styles[s].paragraph_format.space_after = Pt(0)
    section.header.paragraphs[0].text = author
    section.header.paragraphs[0].style = document.styles['Header']
    footer = section.footer.paragraphs[0]
    footer.text = f'{meta.split(" | ")[0]}  |  Seite '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.add_paragraph(title, style='Title')
    document.add_paragraph(meta)
    for block in body.strip().split('\n\n'):
        block = block.strip()
        if block == '---':
            if name not in {'06_Vertragsentwurf_Helix_2026-09-09.docx', '04_Herstellervertrag_DE_2026-08-20.docx'}:
                document.add_page_break()
        elif re.match(r'^\d+(?:\.\d+)*\s', block) and '\n' not in block and len(block) < 110:
            level = 'Heading 2' if '.' in block.split()[0] else 'Heading 1'
            document.add_paragraph(block, style=level)
        else:
            paragraph = document.add_paragraph(block)
            if block.startswith(('Für Wertau Lebensmittel GmbH:', 'Anneke Martens für ', 'Dieser Vertrag wird mit der Unterzeichnung', 'Der Vertrag wird erst durch Unterzeichnung')):
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.keep_together = True
    separate_section_headings(document)
    document.core_properties.author = author
    document.core_properties.title = title
    document.core_properties.subject = meta
    document.core_properties.language = 'en-GB' if '_EN_' in name else 'de-DE'
    document.core_properties.created = datetime(2026, 9, 22, 8, 0)
    document.core_properties.modified = datetime(2026, 9, 22, 8, 0)
    document.save(remember(folder, name))


def mail(folder, name, sender, recipient, date, subject, body, cc=None, reply=None):
    message = EmailMessage(policy=SMTP)
    message['From'], message['To'] = sender, recipient
    if cc:
        message['Cc'] = cc
    message['Date'] = format_datetime(datetime.fromisoformat(date + '+02:00'))
    message['Subject'] = subject
    message['Message-ID'] = f'<{name[:-4]}@{sender.split("@")[1].rstrip(">")}>'
    if reply:
        message['In-Reply-To'] = reply
        message['References'] = reply
    message.set_content(body.strip() + '\n', charset='utf-8')
    message['Content-Language'] = 'en-GB' if body.strip().startswith('Dear ') else 'de-DE'
    remember(folder, name).write_bytes(message.as_bytes())


def txt(folder, name, body):
    remember(folder, name).write_text(body.strip() + '\n', encoding='utf-8')


def csvfile(folder, name, header, rows):
    with remember(folder, name).open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(header)
        writer.writerows(rows)


def project():
    mail(P, '01_2026-08-03_Projektstart.eml', 'Miriam Seidel <m.seidel@wertau-lebensmittel.de>', 'Jonas Rehm <j.rehm@helix-prozessautomation.de>', '2026-08-03T09:12:00', 'Linie 3: Vorplanung und Budget', '''Guten Morgen Herr Rehm,

wir möchten die Dosierung und Verpackung unserer 500-g-Becher in Augsburg zusammenführen. Der vorhandene Becherverschließer und die Etikettierung bleiben stehen. Unser Investitionsrahmen für Ihr Paket liegt bei 675.000 EUR netto. Eine Bestellung des Gesamtprojekts kann ich damit noch nicht auslösen; dafür brauchen wir einen abgestimmten Vertrag und den Beschluss von Frau Dr. Färber.

Die Vorplanung bis zur Layoutfreigabe können wir separat beauftragen. Herr Brandt sammelt die Anschlüsse, Frau Ott aus der Qualitätssicherung die Produktmuster. Bitte rechnen Sie mit Früh- und Spätschicht, auch freitags. Das Umbaufenster im nächsten Sommer ist knapp: Vom 19. bis 30. Juli 2027 steht Linie 3 still, ab 2. August möchten wir wieder produzieren.

Schicken Sie mir bitte zuerst den LOI für die Vorplanung. Die 9.500 EUR hatte ich im Gespräch als Teil des Gesamtbudgets verstanden.

Freundliche Grüße
Miriam Seidel
Leitung Einkauf | Wertau Lebensmittel GmbH
Am Kieswerk 14, 86167 Augsburg''')
    doc(P, '02_LOI_Vorplanung_2026-08-05.docx', 'Wertau Lebensmittel GmbH / Helix Prozessautomation GmbH', 'Vereinbarung über die Vorplanung der Linie 3', 'LOI WA-26-08 | 5. August 2026 | beiderseits freigegebene Fassung', '''1 Beteiligte und Vorhaben

Die Wertau Lebensmittel GmbH, Am Kieswerk 14, 86167 Augsburg, vertreten durch Geschäftsführerin Dr. Eva Färber, beauftragt die Helix Prozessautomation GmbH, An der Werkhalle 8, 86368 Gersthofen, vertreten durch Geschäftsführer Jonas Rehm, mit der Vorplanung einer Dosier- und Verpackungsintegration für Linie 3 in Augsburg. Die bestehende Versiegelungsmaschine und die vorhandene Etikettierung werden nicht ersetzt.

2 Beauftragte Vorplanung

Helix nimmt an einem Aufmaß vor Ort teil, erstellt einen Stellplan mit Wartungsflächen und legt eine Anschlussliste sowie eine erste Taktbetrachtung für 500-g-Becher vor. Eine Laborprobe mit von Wertau bereitgestelltem Produkt gehört zur Vorplanung. Konstruktionszeichnungen zur Fertigung und ausführbare Steuerungssoftware werden in dieser Stufe nicht geschuldet.

Wertau stellt bis 10. August 2026 die verfügbaren Hallenmaße, Schaltpläne und zwei Gebinde Produkt zur Verfügung. Helix übermittelt die Vorplanungsunterlagen bis 4. September 2026. Die Parteien besprechen die Ergebnisse am 8. September 2026. Bleiben Maße ungeklärt, werden diese in der Anschlussliste bezeichnet und beim Ortstermin nachgemessen.

3 Vergütung und Weiterverwendung

Die Vorplanung kostet pauschal 9.500,00 EUR netto zuzüglich 1.805,00 EUR Umsatzsteuer, insgesamt 11.305,00 EUR. Helix stellt nach Unterzeichnung eine gesonderte Rechnung; Wertau zahlt innerhalb von zehn Kalendertagen ab Rechnungseingang. Die Vergütung fällt auch an, wenn kein Hauptvertrag zustande kommt.

Bei Beauftragung der Gesamtanlage wird die Vorplanung im Projektbudget berücksichtigt. Die kaufmännische Zuordnung wird im Hauptvertrag festgehalten. Ein Preis für die Gesamtanlage wird mit dieser Vereinbarung noch nicht festgelegt.

Wertau darf die bezahlten Vorplanungsunterlagen für die Investitionsentscheidung, die Medienplanung und Gespräche mit dem eigenen Hallenplaner verwenden. Die Weitergabe vollständiger Konstruktionsdaten an einen anderen Anlagenbauer bedarf einer gesonderten Verständigung. Einzelne Anschlusswerte darf Wertau an beauftragte Elektro- und Lüftungsunternehmen weitergeben.

4 Weitere Verhandlungen

Keine Partei ist zur Beauftragung oder Annahme des Hauptprojekts verpflichtet. Die Reservierung eines Montagefensters ist unverbindlich, solange der Hauptvertrag nicht geschlossen ist. Änderungen dieses Vorplanungsauftrags werden von den benannten Ansprechpartnern per E-Mail bestätigt, bevor zusätzliche Kosten entstehen.

5 Vertraulichkeit und Abschluss

Produktrezepturen und interne Hallenpläne dürfen ausschließlich zur Bearbeitung des Vorhabens genutzt werden. Helix darf seinem vorgesehenen Kamera-Unterauftragnehmer nur Bildausschnitte ohne Rezepturangaben und die für die Bildaufnahme nötigen Geometriedaten übermitteln. Nach Ende der Gespräche werden nicht benötigte Produktmuster entsorgt und vertrauliche Unterlagen auf Wunsch zurückgegeben; Abrechnungsunterlagen bleiben im Geschäftsarchiv.

Augsburg, 5. August 2026: Dr. Eva Färber, Geschäftsführerin, Wertau Lebensmittel GmbH. Freigabe per E-Mail um 10:06 Uhr.

Gersthofen, 5. August 2026: Jonas Rehm, Geschäftsführer, Helix Prozessautomation GmbH. Gegenbestätigung per E-Mail um 11:18 Uhr.''')
    doc(P, '03_Rechnung_HP-260806.docx', 'Helix Prozessautomation GmbH | An der Werkhalle 8 | 86368 Gersthofen', 'Rechnung zur Vorplanung', 'HP-260806 | Rechnungsdatum 6. August 2026 | Kundenkonto WA-103', '''Rechnungsempfängerin: Wertau Lebensmittel GmbH, Einkauf, Am Kieswerk 14, 86167 Augsburg.

1 Leistung

Vorplanung Dosier- und Verpackungsintegration Linie 3 gemäß LOI WA-26-08 vom 5. August 2026. Pauschale für Aufmaß, Stellplan, Anschlussliste, erste Taktbetrachtung und Laborprobe. Vorgesehener Leistungszeitraum: 6. August bis 8. September 2026.

Menge: 1 Vorplanungspaket. Einzelpreis netto: 9.500,00 EUR. Gesamtbetrag netto: 9.500,00 EUR.

Umsatzsteuer 19 Prozent: 1.805,00 EUR. Rechnungsbetrag brutto: 11.305,00 EUR.

2 Zahlung

Bitte zahlen Sie bis 16. August 2026 ohne Abzug auf die bei Ihnen unter Kreditor 70418 hinterlegte Bankverbindung. Verwenden Sie als Zahlungsreferenz HP-260806 / WA-26-08. Die Rechnung betrifft ausschließlich die Vorplanung; für das Hauptprojekt liegt noch kein Auftrag vor.

Erstellt von Leonie Bach, Rechnungswesen. Elektronische Belegkopie für Einkauf und Buchhaltung. Kontoverbindung und steuerliche Stammdaten werden im freigegebenen Kreditorenstamm geführt.

Leonie Bach
Rechnungswesen | buchhaltung@helix-prozessautomation.de''')
    txt(P, '04_Ueberweisungsbeleg_2026-08-12.txt', '''Wertau Lebensmittel GmbH
Zahlungsverkehr / ausgeführte Einzelzahlung
Exportiert von Jana Moser am 12.08.2026 um 16:42 Uhr

Auftraggeberin: Wertau Lebensmittel GmbH
Belastungskonto: Betriebskonto EUR, Endziffern 4402
Empfängerin: Helix Prozessautomation GmbH
Empfängerkonto: Kreditor 70418, Endziffern 7319
Betrag: 11.305,00 EUR
Verwendungszweck: HP-260806 / WA-26-08 Vorplanung
Auftragsreferenz: WA-ZL-260812-018
Freigegeben: 12.08.2026 11:04 Uhr / Jana Moser und Dr. Eva Färber
Ausgeführt: 12.08.2026 14:18 Uhr
Buchungstag: 12.08.2026
Wertstellung: 12.08.2026
Status: ausgeführt

Zuordnung im Kreditorenbuch: Beleg HP-260806, Ausgleich 11.305,00 EUR.
Anlagenhauptauftrag: noch nicht angelegt.''')
    doc(P, '05_Angebot_HP-2618_2026-09-04.docx', 'Helix Prozessautomation GmbH | Jonas Rehm', 'Angebot für die Integration der Linie 3', 'HP-2618 Rev. 2 | 4. September 2026 | an Miriam Seidel, Wertau Lebensmittel GmbH', '''1 Angebot und Lieferumfang

Wir bieten die Integration der Linie 3 zum Pauschalpreis von 675.000,00 EUR netto an. Das Paket umfasst einen servoelektrischen Zweifachdosierer, die Zuführung ab Übergabepunkt T1, die Fördertechnik bis T4, eine Kamera zur Deckel- und Etikettenkontrolle, einen Ausschleuser, den Linienverteiler sowie die zentrale Ablaufsteuerung. Versiegelungsmaschine, Etikettierer, Hallenverteiler und produktführende Vorbehälter werden von Wertau gestellt.

Die Preisanteile betragen für Mechanik und Dosierung 318.000,00 EUR, Fördertechnik und Schutzumhausung 112.000,00 EUR, Schaltschrank und Elektroinstallation innerhalb der Linie 82.000,00 EUR, Steuerungs- und Kameraeinbindung 96.000,00 EUR sowie Montage, Inbetriebnahme und Einweisung 67.000,00 EUR. Die Summe beträgt 675.000,00 EUR; zuzüglich 128.250,00 EUR Umsatzsteuer ergeben sich 803.250,00 EUR brutto.

2 Ausführung

Die zugesagten Produktdaten beziehen sich auf das Lastenheft LH-WA-03 Rev. 3 und die Helix-Leistungsbeschreibung LB-2618 Rev. 2. Für das 500-g-Produkt sind 90 Becher je Minute am Auslauf T4 vorgesehen. Ob derselbe Takt mit dem zusätzlichen 750-g-Becher erreicht werden kann, ist nicht Gegenstand dieses Angebots. Der hierzu am 2. September geäußerte Wunsch wird separat kalkuliert.

Wir planen mit Vormontage in Gersthofen und einer Werkprüfung vor Versand. Die Einzelmaschinen können vor der Gesamtinbetriebnahme getrennt eingelagert werden. Montagebeginn am 19. Juli 2027 setzt voraus, dass Medien und Hallenboden rechtzeitig bereitstehen. Das Angebot enthält noch keinen Auftrag an das Kameraunternehmen.

3 Vergütung und Zahlungsstufen

Wir schlagen 20 Prozent bei Auftragseingang, 40 Prozent bei Meldung der Versandbereitschaft nach Werkprüfung, 30 Prozent nach mechanischer und elektrischer Fertigstellung vor Ort sowie 10 Prozent nach Abnahme vor. Die jeweiligen Beträge sind 135.000,00 EUR, 270.000,00 EUR, 202.500,00 EUR und 67.500,00 EUR netto. Rechnungen sind innerhalb von 14 Kalendertagen zu zahlen. Die Meilensteine sind noch nicht vereinbart.

Die Vergütung der Vorplanung ist in unserer Projektkalkulation berücksichtigt. Ein zusätzlicher Abzug von 9.500,00 EUR vom Angebotspreis ist in dieser Fassung nicht ausgewiesen. Ihre abweichende Erwartung aus dem Startgespräch haben wir für die Vertragsrunde vorgemerkt.

4 Gültigkeit und Abgrenzung

Das Angebot bleibt bis 30. September 2026 offen. Ein Hauptauftrag kommt erst durch beiderseitige Unterzeichnung des abgestimmten Vertrags zustande. Ein Kran, Stromzuführung zum Linienverteiler, Wasseraufbereitung, Entwässerung und bauliche Durchbrüche sind nicht enthalten. Ein zusätzliches Ersatzteilpaket ist ebenfalls nicht eingerechnet.

Die Konditionen für Abnahme, Softwareüberlassung, Terminverschiebungen und Vertragsbeendigung werden in einem gesonderten Vertragsentwurf vorgeschlagen, den wir am 9. September 2026 übermitteln möchten. Seine Besprechung bleibt vorbehalten.

Jonas Rehm
Geschäftsführer''')
    doc(P, '06_Vertragsentwurf_Helix_2026-09-09.docx', 'Helix Prozessautomation GmbH', 'Vertrag über Lieferung und Integration der Linie 3', 'HP-2618 V1 | 9. September 2026 | Entwurf zur Abstimmung, nicht unterzeichnet', '''1 Vertragsparteien und Unterlagen

1.1 Parteien

Die Helix Prozessautomation GmbH, An der Werkhalle 8, 86368 Gersthofen, vertreten durch Jonas Rehm, übernimmt als Auftragnehmerin die in diesem Vertrag beschriebenen Leistungen für die Wertau Lebensmittel GmbH, Am Kieswerk 14, 86167 Augsburg, vertreten durch Dr. Eva Färber, als Auftraggeberin.

1.2 Vertragsbestandteile

Maßgeblich sind in dieser Reihenfolge dieser Vertrag, das Angebot HP-2618 Rev. 2 vom 4. September 2026, die Leistungsbeschreibung LB-2618 Rev. 2 vom selben Tag, das Schnittstellenbeiblatt SB-03 vom 8. September 2026 und das Lastenheft LH-WA-03 Rev. 3 vom 28. August 2026. Der Stellplan WA-L3-04 vom 8. September 2026 bestimmt die räumliche Anordnung. Bei abweichenden Angaben geht die jeweils zuerst genannte Unterlage vor. Besprechungen verändern den Leistungsumfang nur nach Ziffer 6.

2 Liefergegenstand und Leistung

2.1 Gesamtpaket

Helix liefert den Zweifachdosierer, Förderstrecken, Schutzumhausung, Kamera mit Ausschleusung, Linienverteiler und Ablaufsteuerung. Helix verbindet diese Komponenten am Standort Augsburg mit der vorhandenen Versiegelungsmaschine und dem vorhandenen Etikettierer, nimmt die vereinbarten Funktionen in Betrieb und weist die Bedienmannschaft ein. Die Lieferung gebrauchter Aggregate ist ausgeschlossen. Helix darf fabrikneue Komponenten gleicher Funktion einsetzen, sofern Anschlussmaße und Wartungszugänge unverändert bleiben.

2.2 Bestandsmaschinen

Wertau stellt betriebsbereite Bestandsmaschinen und deren Dokumentation bereit. Eingriffe in deren interne Programme sind nicht geschuldet. Erweist sich eine Bestandsmaschine bei der Einbindung als unzureichend, beschreibt Helix die erforderliche Anpassung und deren Auswirkungen auf den Termin. Die Beauftragung zusätzlicher Arbeiten richtet sich nach Ziffer 6. Die vereinbarte Ausbringung wird nur mit den in der Leistungsbeschreibung bezeichneten Produkten und Bechern erprobt.

2.3 Unterauftragnehmer

Helix darf Optivis Bildsysteme GmbH in Ulm mit Kamerakonfiguration und Bildverarbeitung beauftragen. Wertau richtet technische Hinweise weiterhin an Helix. Helix koordiniert die Leistungen von Optivis und deren Einbindung in die Liniensteuerung. Ein Austausch des Kameraunternehmens wird Wertau vorab mitgeteilt; Wertau kann innerhalb von fünf Arbeitstagen konkrete betriebliche Einwände benennen.

3 Vergütung

3.1 Preis

Der Pauschalpreis beträgt 675.000,00 EUR netto zuzüglich Umsatzsteuer. Die Vorplanung nach LOI WA-26-08 wurde gesondert beauftragt und abgerechnet. Sie wird nicht nochmals berechnet; eine Kürzung des hier vereinbarten Pauschalpreises erfolgt nicht. Leistungen für das 750-g-Format sind erst nach gesonderter Beauftragung Teil des Hauptauftrags.

3.2 Zahlungsfolge

20 Prozent des Nettopreises werden bei Unterzeichnung fällig, 40 Prozent bei schriftlicher Anzeige der Versandbereitschaft nach Durchführung der Werkprüfung, 30 Prozent bei mechanischer und elektrischer Fertigstellung am Standort und 10 Prozent bei Abnahme. Helix legt jeder Rechnung den jeweiligen Fertigstellungsbericht bei. Wertau zahlt innerhalb von 14 Kalendertagen ab Zugang der Rechnung. Einwendungen gegen einen Bericht werden binnen fünf Arbeitstagen mitgeteilt. Die Zahlung stellt für sich keine Abnahme dar.

---

4 Termine und Mitwirkung

4.1 Planungsgrundlage

Die Parteien planen die Werkprüfung für 14. Juni 2027, den Versand für 12. Juli 2027, den Montagebeginn für 19. Juli 2027 und den Leistungslauf am Standort für 29. und 30. Juli 2027. Der Terminplan TP-2618 Rev. 1 setzt die Freigabe der Anschlussdaten bis 18. September 2026 voraus. Die abschließende Terminbestätigung erfolgt binnen zehn Arbeitstagen nach Vertragsschluss. Diese Daten sind bis dahin Planungstermine.

4.2 Bauseitige Voraussetzungen

Wertau stellt eine abgesicherte Zuleitung 400 V, 50 Hz, 125 A bis zum Linienverteiler bereit. Der mögliche Anschlusswert beträgt 72 kW. Druckluft muss am Übergabepunkt bei 6 bar mit mindestens 800 Normlitern pro Minute anliegen. Die Hallenzufahrt ist ab Montagebeginn täglich von 6 bis 18 Uhr frei zu halten. Der Boden im Montagebereich muss trocken, eben und mit 5 kN je Quadratmeter belastbar sein. Wertau stellt einen Elektroverantwortlichen und einen Bediener für die Bestandsmaschinen.

4.3 Fehlende Voraussetzungen

Fehlen notwendige Vorleistungen, meldet Helix dies unter Bezeichnung der betroffenen Arbeit. Die betroffenen Termine verschieben sich um die nachgewiesene Behinderungsdauer zuzüglich der für die Wiederaufnahme erforderlichen Disposition. Helix legt Wertau binnen fünf Arbeitstagen nach Kenntnis einen angepassten Ablauf vor. Bereits eingesetztes Personal wird nur abgerechnet, soweit es nicht anderweitig eingesetzt werden kann; hierfür gelten 960,00 EUR je Monteurstag und 1.180,00 EUR je Inbetriebnehmertag, jeweils netto. Reise- und Unterkunftskosten werden gegen Beleg berechnet.

5 Lieferung und Eigentum

5.1 Transport und Lagerung

Helix organisiert Transport und Abladen innerhalb des Pauschalpreises. Die Gefahr des zufälligen Verlusts der gelieferten Komponenten geht mit der Ablage auf der gekennzeichneten Bereitstellfläche in Augsburg auf Wertau über. Helix trägt die Folgen eigener Montagefehler. Wird die Annahme trotz bestätigter Versandbereitschaft um mehr als zehn Arbeitstage verschoben, darf Helix die Komponenten trocken einlagern und die belegten Lagerkosten weiterberechnen.

5.2 Eigentumsvorbehalt

Die gelieferten Komponenten bleiben bis zur vollständigen Zahlung des Hauptpreises Eigentum von Helix. Wertau darf sie währenddessen zur vertragsgemäßen Produktion nutzen, aber weder veräußern noch zur Sicherung übertragen. Helix benennt die betroffenen Komponenten anhand ihrer Seriennummern. Für von Wertau beigestellte Maschinen wird kein Eigentum von Helix begründet.

6 Änderungen

6.1 Änderungsangebot

Wünscht eine Partei eine Änderung, beschreibt sie die betroffene Funktion, den Anlass und den gewünschten Termin per E-Mail an Miriam Seidel und Jonas Rehm. Helix legt vor Ausführung ein beziffertes Angebot mit Auswirkung auf Ausbringung, Schnittstellen, Dokumentation und Termine vor. Eine technische Besprechung oder die Übersendung eines Musters ersetzt die kaufmännische Freigabe nicht.

6.2 Ausführung

Nur Dr. Eva Färber für Wertau und Jonas Rehm für Helix können kostenpflichtige Änderungen freigeben. Bis zur Freigabe wird nach dem bisherigen Leistungsumfang gearbeitet, soweit dies technisch möglich ist. Ist die Fortführung einzelner Arbeiten wegen einer offenen Änderung nicht sinnvoll, benennt Helix diese Arbeiten und die verbleibenden Alternativen. Eilige Maßnahmen zur Sicherung bereits angelieferter Teile sind zu dokumentieren und Wertau unverzüglich mitzuteilen.

---

7 Prüfungen und Abnahme

7.1 Werkprüfung

Helix lädt Wertau mit zehn Arbeitstagen Vorlauf zur Werkprüfung in Gersthofen ein. Dabei werden Sicherheitskreise, Dosierung, Kamerabild und Schnittstellensignale mit einem Simulator geprüft. Bestandsversiegler und Etikettierer sind im Werk nicht verfügbar. Die Werkprüfung wird mit Messwerten und Abweichungen protokolliert. Sie dient der Versandfreigabe und ersetzt den Leistungslauf der verbundenen Linie in Augsburg nicht.

7.2 Leistungslauf am Standort

Nach Montage und Funktionsprüfung bietet Helix die Gesamtlinie zur Abnahme an. Der Leistungslauf besteht aus vier zusammenhängenden Stunden mit dem 500-g-Produkt bei 6 bis 8 Grad Celsius. Am Auslauf T4 müssen im Mittel mindestens 90 Gutbecher je Minute bei höchstens 1,5 Prozent linienbedingtem Ausschuss erreicht werden. Stillstände von mehr als zwei Minuten wegen fehlenden Produkts, Stroms oder beigestellter Maschinen werden aus der Laufzeit herausgerechnet. Die übrigen Messbedingungen ergeben sich aus LB-2618 Rev. 2. Das Protokoll enthält Start- und Endzeit, Rohzähler sowie die ausgeschlossenen Unterbrechungen.

7.3 Erklärung

Wertau erklärt nach erfolgreichem Lauf die Abnahme schriftlich. Unwesentliche Abweichungen werden mit Beseitigungstermin im Protokoll festgehalten und hindern die Abnahme nicht. Scheitert der Lauf an einer von Helix zu bearbeitenden Abweichung, korrigiert Helix diese und bietet einen Wiederholungslauf an. Die Kosten dieser Wiederholung trägt Helix. Ist die vereinbarte Prüfware nicht verfügbar, vereinbaren die Parteien einen Ersatztermin.

7.4 Nutzung und ausbleibende Erklärung

Nutzt Wertau die Linie nach Abschluss der Inbetriebnahme an zehn Produktionstagen ohne schriftliche Benennung einer wesentlichen Abweichung, gilt sie mit Ablauf des zehnten Tages als abgenommen. Helix weist in der Fertigstellungsanzeige gesondert auf diese Folge hin. Eine vor Abschluss der Inbetriebnahme von Helix begleitete Probefertigung wird nicht auf die zehn Tage angerechnet.

8 Dokumentation und Software

8.1 Übergabe

Helix übergibt vor dem Leistungslauf Bedienungs- und Reinigungsanleitung in deutscher Sprache, Schaltpläne, Stückliste, Ersatzteilliste und Sicherung der projektspezifischen Steuerung. Die abschließend überarbeiteten Unterlagen werden spätestens zehn Arbeitstage nach Abnahme elektronisch geliefert. Die Schnittstellen zur Bestandsanlage und die Not-Halt-Kette sind darin mit dem tatsächlich hergestellten Zustand wiederzugeben.

8.2 Nutzungsumfang

Wertau erhält nach vollständiger Zahlung ein zeitlich unbeschränktes Recht, die gelieferte Software auf Linie 3 zu betreiben und Sicherungskopien anzulegen. Das Steuerungsprojekt wird in bearbeitbarer Form übergeben. Quellcode und Entwicklungsumgebung der Kamerabibliothek von Optivis werden nicht überlassen. Deren Laufzeitlizenz ist für einen Kamerarechner enthalten; die Übertragung auf Ersatzhardware erfordert die Aktivierung durch Optivis.

8.3 Fernzugriff

Fernzugriff ist nur über den von Wertau jeweils freigeschalteten Zugang zulässig. Jede Sitzung wird mit Beginn, Ende und Bearbeiter protokolliert. Produktbilder werden nicht außerhalb der Anlage gespeichert, sofern Wertau nicht im Einzelfall der Übermittlung einer Fehleraufnahme zustimmt. Die Fernwartung wird nach der Sitzung abgeschaltet.

---

9 Mängelbearbeitung und Betreuung

9.1 Meldung und Nacharbeit

Wertau meldet beobachtete Abweichungen unter Angabe von Produkt, Zeit, Fehlermeldung und Zählerstand an Helix. Helix bestätigt den Eingang während der Geschäftszeiten Montag bis Freitag von 8 bis 17 Uhr innerhalb eines Arbeitstages. Helix darf zunächst per Fernzugriff eine Diagnose durchführen. Erforderliche Nacharbeit erfolgt am Standort; ausgetauschte Teile nimmt Helix auf eigene Kosten zurück.

9.2 Zeitraum und Abgrenzung

Helix bearbeitet Mängel der eigenen Leistungen während 24 Monaten ab Abnahme auf eigene Kosten. Schäden durch nicht freigegebene Reinigungsmittel, Umbauten durch Dritte oder Betrieb außerhalb der vereinbarten Produktgrenzen sind davon nicht erfasst, soweit sie auf diesen Umständen beruhen. Verschleißteile werden in der Ersatzteilliste einzeln benannt. Die bloße Aufnahme eines Bauteils in die Liste schließt die Bearbeitung eines ursprünglichen Materialfehlers nicht aus.

9.3 Zusätzlicher Service

Eine Rufbereitschaft außerhalb der Geschäftszeiten ist nicht im Preis enthalten. Wartung, Ersatzteilbevorratung und längere Unterstützung der Kamerasoftware werden separat angeboten. Helix benennt vor Abnahme mindestens einen Ansprechpartner für mechanische Störungen und einen für Steuerung und Kamera. Eine Reaktionszeit bis zur Wiederherstellung des Produktionsbetriebs wird mit diesem Vertrag nicht zugesagt.

10 Verantwortlichkeit und Haftung

10.1 Schäden

Helix haftet unbeschränkt für vorsätzlich oder grob fahrlässig verursachte Schäden sowie für schuldhaft verursachte Verletzungen von Leben, Körper oder Gesundheit. Bei leicht fahrlässiger Verletzung wesentlicher Vertragspflichten wird der Ersatz auf den vorhersehbaren Schaden, insgesamt jedoch auf 675.000,00 EUR begrenzt. Für entgangenen Gewinn, Produktionsausfall und verdorbene Ware wird bei leichter Fahrlässigkeit kein Ersatz geleistet. Zwingende gesetzliche Haftung bleibt unberührt.

10.2 Schadenminderung

Jede Partei informiert die andere unverzüglich über Umstände, die einen größeren Anlagen- oder Produktschaden erwarten lassen. Wertau hält einen erreichbaren Schichtverantwortlichen vor und sichert betroffene Produktchargen. Helix unterstützt die Eingrenzung des technischen Fehlers und erhält die hierfür erforderlichen Betriebsdaten. Ein Anerkenntnis einer Kostenübernahme ist mit dieser Unterstützung nicht verbunden.

11 Unterbrechung und Beendigung

11.1 Abhilfe

Verletzt eine Partei eine wesentliche Pflicht, kann die andere ihr schriftlich eine Frist von 20 Arbeitstagen zur Abhilfe setzen. Die Mitteilung muss die betroffene Pflicht bezeichnen und erkennen lassen, welche Leistung erwartet wird. Nach erfolglosem Ablauf kann die andere Partei den noch nicht erfüllten Teil des Vertrags beenden. Bereits eingetretene Ansprüche werden dadurch nicht aufgehoben.

11.2 Abrechnung bei Projektabbruch

Beendet Wertau das Projekt ohne eine von Helix zu vertretende Pflichtverletzung, vergütet Wertau die bis dahin erbrachten Leistungen sowie nicht stornierbare, projektspezifisch bestellte Teile. Für den entfallenden Rest der Leistung erhält Helix zehn Prozent des darauf entfallenden Nettopreises. Helix legt Leistungsstände, Bestellbelege, Stornierungsmöglichkeiten und bereits geleistete Zahlungen offen; Doppelansätze werden ausgeschlossen.

---

11.3 Herausgabe und Weiterführung

Nach Abrechnung übergibt Helix die bezahlten Komponenten, die vorhandenen projektspezifischen Zeichnungen und den zuletzt gesicherten Stand des Steuerungsprojekts. Unfertige Funktionen werden in einer Übergabeliste bezeichnet. Die Kamerabibliothek wird nur im Umfang der erworbenen Laufzeitlizenz überlassen. Ein Zutritt von Helix zum Werk zur Demontage oder Abholung bedarf eines abgestimmten Termins. Wertau sichert bis dahin die gelagerten Teile gegen Zugriff Unbefugter.

11.4 Längere Unterbrechung

Dauert eine von keiner Partei beherrschbare Unterbrechung länger als 60 Kalendertage, verhandeln die Parteien über Fortsetzung und Termine. Kommt binnen weiteren 20 Kalendertagen keine Verständigung zustande, kann jede Partei die noch ausstehenden Leistungen beenden. Abgerechnet werden ausgeführte Leistungen und nachweislich nicht stornierbare Teile; die Pauschale für den entfallenden Rest aus Ziffer 11.2 fällt in diesem Fall nicht an.

12 Vertrauliche Informationen

Die Parteien verwenden Rezepturen, technische Zeichnungen, Betriebsdaten und Preiskalkulationen der jeweils anderen Partei nur zur Durchführung dieses Projekts. Unterauftragnehmer erhalten nur die für ihren Beitrag erforderlichen Daten und werden entsprechend verpflichtet. Veröffentlichungen mit Namen, Logo oder Hallenaufnahmen von Wertau bedürfen einer vorherigen Freigabe durch Miriam Seidel. Helix darf öffentlich verfügbare allgemeine Informationen weiterverwenden, jedoch keine daraus abgeleiteten Angaben zu Produktionsmengen von Wertau veröffentlichen.

13 Ansprechpartner und Erklärungen

Miriam Seidel koordiniert Einkauf und kaufmännische Fragen für Wertau; Jonas Rehm vertritt Helix in Vertragsfragen. Timo Brandt und Paula König koordinieren die Technik. Die technischen Ansprechpartner dürfen Termine für Begehungen und Tests abstimmen, jedoch weder Preise ändern noch auf Ansprüche verzichten. Mitteilungen über Fertigstellung, Zahlungsstufen, Änderungen oder Beendigung werden an die kaufmännischen Ansprechpartner per E-Mail und zusätzlich als unterzeichnetes Dokument übersandt.

14 Recht und Gerichtsstand

Es gilt deutsches Recht. Das Übereinkommen über Verträge über den internationalen Warenkauf wird ausgeschlossen. Für Streitigkeiten aus diesem Vertrag wird Augsburg als Gerichtsstand vereinbart. Vor einer Klage stimmen sich die Geschäftsführer innerhalb von zehn Arbeitstagen nach einem schriftlichen Gesprächsverlangen ab; unaufschiebbare Maßnahmen bleiben möglich.

15 Abschluss und Anlagenstand

Dieser Vertrag wird mit der Unterzeichnung beider Parteien wirksam. Nebenabreden bestehen nicht; Änderungen werden in einem von beiden Parteien bestätigten Nachtrag festgehalten. Die LOI-Vorplanung bleibt hinsichtlich bereits erbrachter Leistungen und ihrer Nutzung bestehen, soweit dieser Vertrag nichts Abweichendes bestimmt. Sollte eine einzelne Regelung nicht durchführbar sein, verständigen sich die Parteien über eine an ihrem wirtschaftlichen Zweck orientierte Ersatzregelung; die übrigen Vereinbarungen bleiben bestehen.

Für Wertau Lebensmittel GmbH: Dr. Eva Färber, Geschäftsführerin. Unterschrift und Datum stehen noch aus.

Für Helix Prozessautomation GmbH: Jonas Rehm, Geschäftsführer. Unterschrift und Datum stehen noch aus.''')
    doc(P, '07_Kundenanmerkungen_2026-09-11.docx', 'Wertau Lebensmittel GmbH | Miriam Seidel / Timo Brandt / Nora Ott', 'Anmerkungen zum Vertragsentwurf HP 2618 V1', 'WA-EK-0911 | 11. September 2026 | an Jonas Rehm', '''1 Vertragsunterlagen und Gesamtfunktion

Zu Ziffer 1.2: Bitte stellen Sie das Lastenheft vor das Angebot und die Leistungsbeschreibung. Wir haben den Vierstundenlauf aus Ihrem Entwurf gelesen; Frau Ott benötigt je einen Lauf in Früh- und Spätschicht einschließlich Produktwechsel. Die Ausbringung ist am Ende der verbundenen Linie zu messen, nicht am Dosierer. Unsere Zahl 90 Gutbecher pro Minute gilt für die fertig verschlossenen und korrekt etikettierten Becher am T4.

Zu Ziffer 2.2: Wir können die Bestandsmaschinen nicht im Leerlauf mit einer bloßen Taktzahl freigeben. Bei der Begehung am 18. August liefen 96 Becher pro Minute; nach 38 Minuten fiel allerdings ein Etikettensensor aus. Helix soll die Signalübergänge selbst prüfen und uns vor Fertigungsfreigabe sagen, ob diese Maschinen den zugesagten Linienlauf ermöglichen. Ein pauschaler Ausschluss sämtlicher Stillstände der Bestandsmaschinen ist uns zu weit.

2 Preis und Zahlung

Zu Ziffer 3.1: Die Vorplanung über 9.500 EUR netto wurde am 12. August bezahlt. Aus dem Satz im LOI zur Berücksichtigung im Projektbudget hatte ich einen Abzug bei der ersten Rate verstanden. Bitte zeigen Sie uns die Behandlung in einer überarbeiteten Zahlungsübersicht. Frau Färber hat 675.000 EUR netto als Gesamtrahmen einschließlich Vorplanung genannt.

Zu Ziffer 3.2: 60 Prozent sollen vor Anlieferung bezahlt sein; das können wir ohne Absicherung nicht freigeben. Wir schlagen 20 Prozent nach Vertrag und Absicherung der Vorauszahlung, 40 Prozent nach Werkprüfung und tatsächlicher Anlieferung, 30 Prozent nach erfolgreichem Leistungslauf und 10 Prozent nach vollständiger Dokumentation und Erledigung der protokollierten Restpunkte vor. Der bisherige Betrag für die erste Stufe soll bis zur Klärung der Vorplanung nicht als unstreitig ausgewiesen werden.

3 Hallenstrom und Termine

Zu Ziffer 4.1: Die Anschlussdaten können wir am 18. September noch nicht abschließend bestätigen. Der Hallenverteiler wird voraussichtlich erst am 9. Juli 2027 energisiert; ursprünglich war der 11. Juni vorgesehen. Das ist eine Verschiebung auf unserer Seite. Bitte prüfen Sie, welche Tests Sie bis dahin in Gersthofen erledigen können und ob die Montage im Juli trotzdem gehalten werden kann.

Zu Ziffer 4.3: Wir benötigen eine tageweise Darstellung, welche Mannschaft wegen welcher fehlenden Vorleistung wartet. Die Wiederanlaufzeit soll nicht ohne Obergrenze zusätzlich zur Wartezeit entstehen. Wir möchten über Standkosten entscheiden können, bevor sie weiterlaufen. Das im Entwurf genannte tägliche Zeitfenster darf die ab 14 Uhr laufende Nachbarlinie nicht blockieren.

---

4 Prüfungen und Abnahme

Zu Ziffer 7.1: Die Werkprüfung mit Simulator ist sinnvoll. Sie darf weder die Funktionsverantwortung für die reale Linie ersetzen noch die letzte Nachweismöglichkeit für die Kamera sein. Mit weißen Testdeckeln wurde das aktuelle Etikett noch nicht geprüft. Bitte verwenden Sie die von Frau Ott markierten Rollen aus Charge E26-082.

Zu Ziffer 7.2: Bitte übernehmen Sie unsere Messregel aus Ziffer 4 des Lastenhefts. Bei nur vier Stunden am Stück sehen wir den Neustart nach der Reinigung nicht. Wir möchten Störungen zunächst vollständig im Protokoll behalten; welche Zeiten herausgerechnet werden, sollen die Verantwortlichen gemeinsam festhalten. Ein Stromausfall ist etwas anderes als eine von der Liniensteuerung nicht abgefangene kurze Pause des Versieglers.

Zu Ziffer 7.4: Die Nutzung an zehn Tagen kann in unserem Anlaufbetrieb nötig sein, obwohl Restpunkte offen sind. Wir wollen eine ausdrückliche Abnahmeerklärung und keine Abnahme allein durch Produktion. Frau Ott unterschreibt Messprotokolle, aber nicht die Abnahme; dafür soll Frau Färber zeichnen. Helix soll einen wiederholten Lauf nach eigener Nacharbeit nicht gesondert berechnen.

5 Software und Betreuung

Zu Ziffer 8.2: Unsere Instandhaltung muss die Anlage auch nach Ausfall des Kamerarechners hochfahren können. Herr Brandt benötigt einen getesteten Rücksicherungsweg, Lizenzdaten, Administratorzugang und eine Ersatzhardwareliste. Bitte erläutern Sie, wer aktiviert, wenn Optivis nicht mehr erreichbar ist. Ein bearbeitbares Steuerungsprojekt ohne lauffähige Kameralizenz genügt dafür nicht.

Zu Ziffer 9.3: Zwischen 14 und 22 Uhr läuft die zweite Schicht. Eine Rückmeldung erst am nächsten Arbeitstag hilft bei einem stillstehenden Ausschleuser wenig. Wir bitten um ein gesondertes Angebot für telefonische Unterstützung bis 22 Uhr und einen benannten Ersatzansprechpartner. Das darf gern separat bepreist werden; wir wollen vor Abschluss wissen, was es kostet.

6 Schäden und vorzeitiges Ende

Zu Ziffer 10.1: In einem Produktionslauf stehen etwa 6.000 EUR Produktwert vor der Verpackung. Eine verspätete Ausschleusung kann außerdem einen bereits etikettierten Bestand betreffen. Den vollständigen Ausschluss verdorbener Ware und Produktionsausfall akzeptiert Frau Färber nicht. Bitte lassen Sie Ihren Versicherer die mit unserem Produkt verbundenen Positionen erläutern und schicken Sie eine Deckungsbestätigung.

Zu Ziffer 11.2: Wir möchten bei einem Abbruch die Bestellbelege sehen und selbst entscheiden können, ob wir angefangene Baugruppen übernehmen. Zehn Prozent für den gesamten nicht ausgeführten Rest plus nicht stornierbare Teile kann ich nicht freigeben, ohne zu wissen, was im Rest noch enthalten ist. Zu Ziffer 11.3 benötigen wir ein Recht, die bezahlten Planungsstände durch ein anderes Integrationsunternehmen fertigstellen zu lassen. Für einen Abbruch wegen ausbleibender Helix-Leistung fehlt uns eine konkrete Übergabefrist.

Miriam Seidel
Mit Timo Brandt und Nora Ott am 11. September abgestimmt; noch keine Freigabe des Gesamtvertrags.''')
    doc(P, '08_Lastenheft_LH-WA-03_Rev3.docx', 'Wertau Lebensmittel GmbH | Technik und Qualitätssicherung', 'Lastenheft für Dosierung und Verpackung Linie 3', 'LH-WA-03 Rev. 3 | 28. August 2026 | Nora Ott / Timo Brandt', '''1 Produkt und Betriebsweise

Linie 3 verarbeitet einen gekühlten, stückigen Frischkäseaufstrich. Die Produkttemperatur am Dosierer liegt zwischen 6 und 8 Grad Celsius. Die Dichte der Grundmasse liegt bei 1,04 bis 1,08 kg je Liter. Kräuter und Gemüsestücke erreichen eine Kantenlänge von 6 mm. Der Dosierer darf diese Stücke nicht zerreiben. In den Mustern vom 12. August beträgt der Feststoffanteil 8 bis 11 Prozent nach Masse.

Das Nennfüllgewicht beträgt 500 g je Becher. Der Becher hat oben 95 mm Außendurchmesser, ist 92 mm hoch und wird auf einem 110-mm-Raster zugeführt. Pro Schicht werden bis zu 32.000 Becher gefertigt. Die Frühschicht läuft von 6 bis 14 Uhr, die Spätschicht von 14 bis 22 Uhr, jeweils mit 30 Minuten geplanter Pause. Zwischen 22 und 23:30 Uhr wird gereinigt. An fünf Tagen je Woche können beide Schichten belegt sein.

Der 750-g-Becher ist 125 mm hoch bei gleichem Deckeldurchmesser. Er wird für Herbst 2027 erwogen, ist aber noch nicht als Serienformat freigegeben. Die Halterung soll nach Möglichkeit ausreichend Verstellweg bieten; eine Produktionsfreigabe für dieses Format ist nicht Gegenstand der aktuellen Musterprobe.

2 Hygiene und Reinigung

Produktberührende Teile bestehen aus Edelstahl 1.4404 oder dem vorab von Nora Ott freigegebenen Dichtungswerkstoff. Toträume an Ventilen sollen beim Zerlegen zugänglich sein. Werkzeuglos zu entnehmende Teile dürfen höchstens 12 kg wiegen; schwerere Teile benötigen eine gekennzeichnete Hebeaufnahme. Tropfstellen über offenen Bechern sind zu vermeiden. Schraubgewinde dürfen nicht in den Produktweg ragen.

Die tägliche Reinigung beginnt mit Entleerung und Vorspülung bei 35 bis 40 Grad Celsius, anschließend werden Dosierköpfe zerlegt. Für die alkalische Reinigung sind 1,5 Prozent Reiniger bei höchstens 55 Grad Celsius und 20 Minuten Einwirkzeit vorgesehen. Danach folgen Klarspülung und Sichtkontrolle. Elektrische Baugruppen werden nicht direkt mit dem Schlauch abgespritzt. Helix bezeichnet die hierzu zu schützenden Bereiche dauerhaft an der Anlage.

Vom Produktionsstopp bis zur freigegebenen Sichtkontrolle sollen zwei eingewiesene Personen höchstens 75 Minuten benötigen. Die Kameraabdeckung muss ohne Neueinrichtung der Kamera abnehmbar sein. Nach Wiedereinbau darf die Referenzlage höchstens 0,5 mm vom vorherigen Zustand abweichen. Wertau dokumentiert die Reinigung am vorgesehenen Testtag mit Beginn, Ende und Befund.

---

3 Ausbringung und Qualität

Bei 500 g Nennfüllgewicht werden am T4 mindestens 90 verkaufsfähige Becher je Minute verlangt. Ein verkaufsfähiger Becher ist verschlossen, richtig etikettiert, innerhalb des vereinbarten Gewichtsfensters und ohne erkennbares Produkt am Siegelrand. Das Gewichtsfenster für den Leistungslauf beträgt 495 bis 507 g; der Mittelwert aus jeweils 30 Bechern muss zwischen 500 und 503 g liegen. Tara wird vor jedem Stichprobenblock aus zehn leeren Bechern ermittelt.

Der linienbedingte Ausschuss soll höchstens 1,5 Prozent der zugeführten Becher betragen. Die Kamera prüft Deckelanwesenheit, Etikettenanwesenheit, Druckbild des Chargencodes und die Zuordnung der aufgerufenen Rezeptnummer. Für die Prüfung werden je 20 absichtlich falsch vorbereitete Becher pro Fehlerart markiert und eingeschleust. Alle markierten Fehlbecher müssen in der gesicherten Ausschussbox landen. Unter 500 ordnungsgemäß vorbereiteten Kontrollbechern sollen höchstens fünf fälschlich ausgeschleust werden.

Eine gefüllte Ausschussbox muss den Zufluss stoppen, bevor ein weiterer Fehlbecher den Ausschleuspunkt passiert. Die Quittierung darf nur nach Entleerung vor Ort erfolgen. Bei Ausfall der Kamera darf die Anlage nicht unbemerkt in den normalen Gutlauf wechseln. Ein ausdrücklich freigegebener Handbetrieb muss am Bedienpult sichtbar sein und die betroffene Charge im Protokoll kennzeichnen.

4 Leistungslauf am Standort

Wertau erwartet einen vierstündigen Lauf in der Frühschicht und einen vierstündigen Lauf in der Spätschicht am Folgetag, mit Reinigung und Neustart dazwischen. Pro Stunde werden 30 aufeinanderfolgende Becher gewogen. Der Auslaufzähler wird zu Beginn und Ende sowie nach jedem Stopp abgelesen. Alle Stopps ab zehn Sekunden werden mit Uhrzeit, Dauer und beobachtetem Anlass erfasst. Die Rohdaten bleiben dem Protokoll beigefügt.

Produktmangel und nachweisliche externe Stromunterbrechungen werden gesondert ausgewiesen. Die Anrechnung oder Herausrechnung anderer Stopps wird vor Unterzeichnung des Protokolls gemeinsam vermerkt; ein automatisches Streichen von Stillständen der Bestandsmaschinen ist nicht vorgesehen. Ein Fehler der Not-Halt-Kette, ein nicht ausgeschleuster markierter Fehlbecher oder ein offener Schutzkreis beendet den jeweiligen Lauf. Der Wiederholungslauf beginnt nach Beseitigung neu.

5 Bedienung und Unterlagen

Die Rezeptumschaltung erfolgt am Bedienpult durch Schichtleiter. Bediener dürfen Sollwerte nur innerhalb festgelegter Grenzen verändern. Reinigung und Instandhaltung erhalten eigene Zugänge. Wertau benötigt Schaltplan, Ersatzteilliste, Reinigungsanweisung, Sicherungen der Software und die Zuordnung der Zugangsrollen spätestens zum Leistungslauf. Vier Schichtleiter und zwei Instandhalter sind getrennt einzuweisen.''')
    doc(P, '09_Leistungsbeschreibung_LB-2618.docx', 'Helix Prozessautomation GmbH | Paula König, Projektleitung', 'Leistungsbeschreibung Linie 3', 'LB-2618 Rev. 2 | 4. September 2026 | Bezug Angebot HP-2618', '''1 Mechanischer Lieferumfang

Der Zweifachdosierer HP-D2 hat einen fahrbaren Edelstahlrahmen mit zwei synchronisierten Dosierachsen und je 650 ml maximalem Hubvolumen. Für das 500-g-Produkt werden zwei Becher pro Hub beschickt. Eine geregelte Vorschubstrecke übernimmt die leeren Becher am T1. Die bestehende Versiegelungsmaschine beginnt am T2. Helix liefert die verbindenden Bänder und eine Staustrecke für 28 Becher vor dem Versiegler. Der Bestandsetikettierer folgt nach dem Versiegler; Kamera und Ausschleuser liegen zwischen T3 und T4.

Die Gesamtfläche der neuen Baugruppen beträgt 8,40 m in Förderrichtung und 2,40 m in der Tiefe einschließlich Umhausung. Der Bediengang auf der Südseite bleibt mit 1,20 m frei. Türen der Umhausung werden nach Norden geöffnet; dort ist ein Wartungsraum von 0,90 m erforderlich. Die Schaltschranktür liegt außerhalb des nassen Reinigungsbereichs. Förderhöhe ist 920 mm, einstellbar um 30 mm.

2 Steuerung und Kamera

Helix liefert die Liniensteuerung mit 12-Zoll-Bedienpult, Rezepthandhabung für das 500-g-Produkt und Ereignisspeicher für 90 Tage. Der Speicher erfasst Stoppbeginn, Stoppende, auslösenden Kanal, aktives Rezept und Zählerstand. Wertau erhält einen CSV-Export über einen lokalen USB-Anschluss. Eine Verbindung zur Produktionsdatenbank von Wertau ist nicht enthalten.

Optivis liefert eine fest montierte Farbkamera, Beleuchtung, Industrie-PC und Laufzeitlizenz der Bildauswertung. Das Bildfenster misst 110 mal 110 mm. Die Kamera prüft Deckel und Etikett von oben; das seitliche Druckbild wird über einen Spiegel erfasst. Ein vollständiges Lesen verschmutzter oder geknitterter Etiketten wird nicht zugesagt. Die eingereichten Originalrollen werden bis zur Werkprüfung zur Parametrierung verwendet.

3 Montage und Einweisung

Der Preis enthält fünf Tage mechanische Montage mit drei Personen, fünf Tage Elektroinstallation mit zwei Personen und zehn Tage Inbetriebnahme mit einem Steuerungstechniker. Diese Ansätze begrenzen nicht die geschuldete Fertigstellung des vereinbarten Pakets. Vom Kunden verursachte Wartezeiten werden getrennt dokumentiert. Zwei Einweisungen von je vier Stunden sind vorgesehen; ein Termin wird in der Spätschicht durchgeführt.

---

4 Werkprüfung und Messbedingungen

Die Werkprüfung dauert einen Arbeitstag. Der Dosierer läuft mit 120 kg Wertau-Produkt. Die Bestandsmaschinen werden durch Signalsimulatoren ersetzt. Geprüft werden Rezeptwechsel, Stillsetzung, Wiederanlauf, Ausfall der Druckluft und 80 markierte Kamerafehler. Gewichte aus mindestens zwölf Stichproben werden elektronisch übergeben. Diese Prüfung liefert noch keine Ausbringungszahl der verbundenen Bestandslinie.

Am Standort wird ein vierstündiger Dauerlauf vorgeschlagen. Ausbringung und Ausschuss werden am T4 aus den Rohzählern bestimmt. Bei 90 Gutbechern je Minute sind während vier Stunden ungekürzter Messzeit 21.600 Gutbecher erforderlich. Der Vorschlag erfasst keinen zweiten Lauf nach Reinigung. Dieser Punkt sowie die Behandlung von Bestandsmaschinenstopps weichen vom Lastenheft Rev. 3 ab und wurden noch nicht mit Wertau abgestimmt.

5 Liefergrenzen

Helix stellt Leitungen und Druckluftschläuche ab den Übergabepunkten des Schnittstellenbeiblatts. Eine neue zentrale Stromzuleitung, zusätzliche Entwässerung, die Anpassung des Hallenbrandschutzes und ein Umbau des internen Versieglerprogramms gehören nicht zum Preis. Auf Produktleitungen vor dem Dosierer wird nur nach gesonderter Freigabe zugegriffen. Die von Wertau bereitzustellende Pumpe muss 50 Liter je Minute ohne Lufteintrag liefern.

6 Unterlagen und Rücksicherung

Das Steuerungsprojekt wird mit Versionsliste, Parametern und lesbarem Änderungsprotokoll übergeben. Die Kamera wird als lauffähiges Rechnerabbild gesichert. Zur Aktivierung auf einem Ersatzrechner muss Optivis eine neue hardwaregebundene Lizenzdatei ausstellen. Helix übernimmt den Aktivierungsantrag während der Inbetriebnahme. Eine unabhängig von Optivis nutzbare Lizenzreserve ist in der bisherigen Kalkulation nicht enthalten.

Paula König
Projektleitung Helix Prozessautomation GmbH''')
    doc(P, '10_Schnittstellenbeiblatt_SB-03.docx', 'Helix Prozessautomation GmbH / Wertau Lebensmittel GmbH', 'Schnittstellenbeiblatt Linie 3', 'SB-03 | 8. September 2026 | technische Besprechungsfassung', '''1 Medien und Raum

Der elektrische Übergabepunkt E1 liegt im nordwestlichen Anschlussbereich bei x = 1,0 m und y = 5,5 m gemäß Stellplan WA-L3-04. Wertau führt 400 V, 50 Hz und 125 A bis zur Anschlussklemme im Helix-Linienverteiler. Der voraussichtliche Anschlusswert ist 72 kW. Helix installiert ab dieser Klemme und prüft die Schutzleiterverbindung innerhalb des neuen Pakets. Das Messprotokoll der vorgeschalteten Zuleitung stellt Wertau vor dem ersten Einschalten bereit.

Druckluft wird am Punkt D1 bei x = 3,0 m und y = 5,5 m mit 6 bar Fließdruck übergeben. Der Bedarf beträgt 800 Normliter je Minute, kurzzeitig 1.100 Normliter während gleichzeitiger Ausschleusung und Ventilreinigung. Am 18. August wurden bei laufender Nachbarlinie 5,7 bar gemessen. Timo Brandt prüft einen zusätzlichen Pufferspeicher; dessen Beschaffung ist nicht freigegeben.

Am Produktübergang P1 wird der Dosierer mit einem DN-40-Anschluss versorgt. Wertau stellt 6 bis 8 Grad Celsius Produkttemperatur und eine Pumpe mit mindestens 50 Litern je Minute bereit. Die Entfernung vom Vorbehälter bis P1 beträgt 3,8 m. Der Reinigungsauslauf R1 liegt bei x = 2,2 m und y = 0,5 m; Wertau hält den Bodenablauf frei.

2 Maschinensignale

Zwischen Dosierer und Versiegler werden die Signale Betriebsbereit, Freigabe Zulauf, Störung und Not-Halt zweikanalig beziehungsweise für Betriebsdaten als 24-V-Signale übergeben. Der Versiegler meldet die Freigabe für jeden Becher. Helix puffert höchstens 28 Becher. Bleibt die Freigabe länger als 800 ms aus, stoppt der Dosierer; die Wiederfreigabe wird mit einer Verzögerung von 2 s übernommen.

Der Etikettierer erhält die Rezeptnummer 500 und die Charge als ASCII-Zeichenfolge mit höchstens zwölf Zeichen. Timo Brandt hat bislang nur eine Liste der Ein- und Ausgänge, aber keine aktuelle Programmsicherung gefunden. Der Etikettierer meldet Etikett vorhanden an die Linie, während die Kamera die tatsächliche Anwesenheit am Becher nochmals prüft. Das Signal Ausschleusen wird um die gemessene Bandlaufzeit von 740 ms verzögert.

3 Netzwerk und Daten

Die neuen Steuerungen liegen im abgeschotteten Maschinennetz 192.168.30.0/24. Wertau reserviert für den Kamera-PC die Adresse 192.168.30.20. Der Linienrechner führt keinen Internetzugang. Fernwartung erfolgt über einen von Wertau bereitgestellten Zugang, der durch den Schichtverantwortlichen zeitlich freigeschaltet wird. Ein Export enthält Zeitstempel mit Zeitzone, Rezeptnummer, Gut- und Ausschusszähler sowie Störungscode. Bilddaten werden nach 30 Tagen lokal überschrieben, ausgenommen manuell gesicherte Fehlerbilder.

4 Mechanische Übergänge

T1 liegt vor dem Dosierer bei x = 0,8 m, T2 am Versieglereinlauf bei x = 4,2 m, T3 nach dem Etikettierer bei x = 6,3 m und T4 nach dem Ausschleuser bei x = 8,8 m. Alle Übergänge liegen bei y = 2,6 m und 920 mm Förderhöhe. Die freie Übergabelücke darf höchstens 8 mm betragen. Für den Bechertransport werden seitliche Führungen mit 3 mm Gesamtspiel eingestellt.

5 Noch zu bestätigende Anschlussdaten

Wertau hat die Energieversorgung E1 noch nicht zur Ausführung freigegeben. Die Nachricht des Elektroplaners vom 10. September nennt den 9. Juli 2027 als frühesten Zuschalttermin. Paula König benötigt bis 18. September 2026 die Entscheidung, ob der bisherige Montageplan beibehalten werden soll. Außerdem fehlen das aktuelle Programm des Etikettierers und die Bestätigung des Druckluftpuffers. Die im Stellplan eingetragene Bereitstellfläche wird derzeit von Paletten der Nachbarlinie genutzt.

Für Wertau aufgenommen: Timo Brandt, Instandhaltung.
Für Helix aufgenommen: Paula König, Projektleitung.
Die technische Besprechung stellt noch keine kaufmännische Freigabe dar.''')
    doc(P, '11_Terminplan_TP-2618_Rev1.docx', 'Helix Prozessautomation GmbH | Paula König', 'Terminplan für Linie 3', 'TP-2618 Rev. 1 | 9. September 2026 | Planungsstand ohne Hauptauftrag', '''1 Vorbereitung im Jahr 2026

Bis 18. September 2026 sollen Anschlussdaten, Reinigungsvorgaben und Becherzeichnungen bestätigt werden. Der Vertragsabschluss wird für 30. September 2026 angestrebt. Die mechanische Konstruktion ist vom 5. Oktober bis 18. Dezember 2026 vorgesehen. Vor Bestellung der Dosierachsen erhält Wertau den Einbauplan mit Wartungsflächen zur Freigabe. Helix benötigt eine Rückmeldung innerhalb von fünf Arbeitstagen, weil die Achsen mit 18 Wochen Lieferzeit angeboten wurden.

2 Fertigung und Werkprüfung im Jahr 2027

Die Baugruppen sollen vom 11. Januar bis 28. Mai 2027 gefertigt werden. Die Softwarevorbereitung läuft vom 1. März bis 4. Juni 2027 mit dem Signalsimulator. Optivis hat vorläufig die Kalenderwoche ab 31. Mai 2027 für die Parametrierung reserviert; eine verbindliche Buchung ist noch nicht erfolgt. Die Werkprüfung in Gersthofen ist am 14. Juni 2027 vorgesehen. Wertau stellt bis 7. Juni 120 kg Produkt, 2.000 Becher und 2.000 Originaletiketten bereit.

3 Montagefenster Augsburg

Der bisherige Ablauf setzt bauseitige Energiebereitschaft am 11. Juni 2027 voraus. Die Lieferung soll am 12. Juli erfolgen; bis dahin benötigt Helix die Bereitstellfläche Nord mit 18 Quadratmetern. Die bestehende Linie wird am 16. Juli nach der Spätschicht außer Betrieb genommen. Vom 19. bis 23. Juli erfolgt die mechanische Montage, vom 26. bis 28. Juli die verbundene Inbetriebnahme. Am 29. und 30. Juli sind Prüf- und Reservezeiten eingeplant. Wertau möchte am 2. August wieder regulär produzieren.

Der vorstehende Ablauf enthält den von Helix vorgeschlagenen Vierstundenlauf und einen Reservetag. Die zwei von Wertau gewünschten Läufe mit Reinigung zwischen Früh- und Spätschicht sind darin noch nicht als feste Abnahmetermine abgebildet. Ein zusätzlicher Lauf kann den Reservetag verbrauchen.

4 Voraussetzungen und Disposition

Für die Vormontage werden keine Betriebsdaten aus Augsburg benötigt; für die Signalprüfung am Standort braucht Helix den betriebsfähigen Hallenanschluss und die Programme der Bestandsmaschinen. Wird die Zuschaltung von Strom nach hinten verschoben, prüft Paula König zunächst, welche Prüfschritte vorgezogen werden können. Ein verbindlicher Ersatztermin wird erst nach Rückmeldung der Elektroplanung und von Optivis vorgeschlagen.

Dieser Plan ersetzt keine Terminbestätigung nach Vertragsschluss. Der Zahlungsplan in der Arbeitsmappe verwendet dieselben Termine lediglich als Liquiditätsannahmen; eine bereits entstandene Zahlungsforderung wird damit nicht ausgewiesen.''')
    mail(P, '12_2026-09-10_Stromversorgung.eml', 'Timo Brandt <t.brandt@wertau-lebensmittel.de>', 'Paula König <p.koenig@helix-prozessautomation.de>', '2026-09-10T16:24:00', 'E1: Zuschaltung erst am 09.07.2027 möglich', '''Hallo Frau König,

ich habe heute die Rückmeldung unserer Elektroplanung erhalten. Der neue Hallenverteiler kommt nicht bis 11. Juni. Nach dem abgestimmten Bauablauf können wir E1 frühestens am 9. Juli 2027 zuschalten. Das betrifft unsere Hallenversorgung, nicht Ihren Schaltschrank. Ein 32-A-Baustromanschluss ist vorhanden, für den vollständigen Probelauf reicht er nicht.

Den Einzug der Zuleitung können wir im Juni vorbereiten. Eine Prüfung unter Last ist vorher aber nicht möglich. Unser Produktionsleiter hält am Stillstand 19. bis 30. Juli fest; einen dritten Stillstandsblock haben wir bislang nicht vorgesehen.

Bitte sagen Sie uns, ob Ihre Techniker direkt am 19. Juli anfangen können oder ob Sie vorher noch einmal unter Spannung prüfen müssen. Die Palette vor dem Nordtor steht noch dort. Ich kläre den Räumtermin mit der Logistik.

Viele Grüße
Timo Brandt
Instandhaltung | Wertau Lebensmittel GmbH
Am Kieswerk 14, 86167 Augsburg''', cc='Miriam Seidel <m.seidel@wertau-lebensmittel.de>')
    mail(P, '13_2026-09-14_Terminantwort.eml', 'Paula König <p.koenig@helix-prozessautomation.de>', 'Timo Brandt <t.brandt@wertau-lebensmittel.de>', '2026-09-14T08:35:00', 'AW: E1 und unser Julifenster', '''Hallo Herr Brandt,

mit dem 9. Juli entfällt unser bisher vorgesehener Vorabtermin am angeschlossenen Verteiler. Wir könnten die Signalprüfung mit Ihrem Etikettierer auf den 12. Juli legen, wenn die Freigabe der Zuleitung dann vorliegt. Für diesen Tag benötigen wir Zugang zu beiden Bestandsmaschinen und eine Person, die sie bedienen kann.

Optivis hat den Kameratechniker bisher vom 26. bis 28. Juli vorgemerkt. Er ist ab 2. August auf einer anderen Baustelle. Wenn die Linie am 26. Juli noch nicht elektrisch durchgängig läuft, bekommen wir voraussichtlich erst ab 16. August wieder jemanden. Das ist noch keine bestätigte Verschiebung; ich möchte nur die aktuelle Belegung offenlegen.

Ich sehe eine Chance für den bisherigen Juliablauf, aber keinen belastbaren Puffer für einen erneuten Aufbau nach der Reinigung. Den Zweischichtnachweis von Frau Ott müssen wir gemeinsam einplanen. Standkosten sind bislang nicht entstanden, wir haben noch kein Personal bestellt.

Freundliche Grüße
Paula König
Projektleitung | Helix Prozessautomation GmbH
An der Werkhalle 8, 86368 Gersthofen''', cc='Miriam Seidel <m.seidel@wertau-lebensmittel.de>; Jonas Rehm <j.rehm@helix-prozessautomation.de>')
    doc(P, '14_Optivis_Integrationsangebot_OV-144.docx', 'Optivis Bildsysteme GmbH | Kellerstraße 18 | 89077 Ulm', 'Kameraeinbindung für Helix Linie WA 3', 'OV-144 | 15. September 2026 | an Paula König', '''1 Paket und Preis

Wir bieten Helix Prozessautomation GmbH die Deckel-, Etiketten- und Codekontrolle der Linie WA-3 für 38.400,00 EUR netto an. Enthalten sind Kamera, Beleuchtung, Spiegeloptik, Industrie-PC, eine Laufzeitlizenz, Parametrierung für das 500-g-Format und drei Einsatztage in Augsburg. Helix stellt mechanische Halter, Schutzabdeckung, Trigger, Ausschleussignal und Zugang zur Liniensteuerung. Dieses Angebot ist Bestandteil Ihrer Integrationskalkulation und kein zusätzliches Angebot an Wertau.

2 Erkennung und Signalweg

Wir werten Deckel und Etikett bei konstantem Bandlauf bis 18 m je Minute aus. Die Auswertung liefert binnen 120 ms nach Trigger entweder Gut, Ausschleusen oder Bild ungültig. Helix muss auch bei Bild ungültig ausschleusen und den Status protokollieren. Die Freigabe des Chargencodes erfolgt anhand der von Wertau bereitgestellten Zeichenfolge. Das Durchlesen stark reflektierender Knickstellen kann einen zweiten Beleuchtungswinkel erfordern; dessen Erprobung gehört zur Werkparametrierung.

3 Software und Ersatzrechner

Die Laufzeitlizenz ist an die Kennung des Industrie-PC gebunden. Ein Rechnerabbild allein aktiviert die Auswertung nicht auf anderer Hardware. Eine Ersatzaktivierung erfolgt nach Angabe der neuen Gerätekennung durch unseren Service. Während der ersten 24 Monate ab Standortinbetriebnahme ist eine solche Aktivierung enthalten; anschließend berechnen wir 480,00 EUR netto je Vorgang. Ein zweiter vollständig lizenzierter Rechner ist optional für 4.900,00 EUR netto erhältlich und in den 38.400,00 EUR nicht enthalten.

Der Quellcode der Bildbibliothek wird nicht herausgegeben. Projektparameter, Prüfrezepte und Ergebnisprotokolle darf Helix an Wertau weitergeben. Wertau darf die Laufzeit auf genau einer Linie nutzen. Für eine Weiterführung durch einen anderen Integrator müssen wir dessen Zugriff auf unsere Konfigurationsoberfläche gesondert freischalten.

4 Termine und Support

Die Werkparametrierung ist für 31. Mai bis 2. Juni 2027 reserviert. Der Standorttermin 26. bis 28. Juli 2027 wird bis 25. September 2026 vorgehalten. Danach ist eine neue Bestätigung erforderlich. Unser nächster derzeit freier Standorttermin beginnt am 16. August 2027. Telefonische Unterstützung ist montags bis freitags zwischen 8 und 17 Uhr eingeschlossen; eine Unterstützung bis 22 Uhr wird nach Personalabstimmung gesondert angeboten.

5 Abrechnung und Bestätigung

40 Prozent werden nach schriftlicher Bestellung, 40 Prozent nach Lieferung der Hardware und 20 Prozent nach Standortinbetriebnahme mit 14 Tagen Zahlungsziel berechnet. Die Lieferzeit beginnt nach Eingang der Originaletiketten und Becherzeichnungen. Änderungen des 750-g-Formats sind nicht enthalten. Dieses Angebot gilt bis 25. September 2026 und ist noch nicht bestellt.

Elena Haas
Vertrieb und Anwendungstechnik | e.haas@optivis-bildsysteme.de''')
    txt(P, '16_Chat_Baufreiheit_2026-09-16.txt', '''Projektchat Linie 3
Export: Timo Brandt, 16.09.2026 15:41 Uhr
Teilnehmer: Timo Brandt (Wertau), Paula König (Helix), Sven Albers (Wertau Logistik)

16.09.2026 14:07 | Paula König
Ist die Nordfläche wirklich ab 12.07. frei? Im Plan sind 18 Quadratmeter für die Anlieferung eingetragen. Auf dem Foto vom Ortstermin standen dort noch Folienrollen.

16.09.2026 14:11 | Sven Albers
Die Rollen können weg. Die gelben Paletten nicht vor dem 16.07. Die gehören zum letzten Auftrag vor dem Stillstand. Wir brauchen davon täglich Material.

16.09.2026 14:14 | Timo Brandt
Ich hatte in der Besprechung gesagt, wir räumen bis Montag. Gemeint war Montag, 19.07., nicht der Anliefertag. Das war wohl missverständlich.

16.09.2026 14:19 | Paula König
Dann können wir am 12.07. zwar an die Steuerung, aber nicht abladen? Der Lkw ist noch nicht gebucht. Ich brauche 3 mal 6 Meter trocken und einen freien Weg mit 2,4 Metern Breite vom Tor.

16.09.2026 14:24 | Sven Albers
Westlager ginge vielleicht. Das Tor dort hat 2,1 Meter. Könnt ihr die Umhausung erst hier montieren?

16.09.2026 14:29 | Paula König
Die größte Baugruppe ist 2,25 Meter breit. Durch 2,1 Meter passt sie nicht. Zerlegen müssten wir neu planen. Ich lasse die Anlieferung bis zur Klärung auf dem 12.07., aber nicht als bestätigt.

16.09.2026 14:35 | Timo Brandt
Ich spreche morgen mit der Schichtleitung. Bitte noch keinen zweiten Transport einplanen. Strom am 09.07. ist ebenfalls nur der früheste Termin, das Prüfprotokoll habe ich noch nicht.

16.09.2026 15:03 | Sven Albers
Für die Montage ab 19.07. bekomme ich die Fläche frei. Vorher kann ich es heute nicht zusagen.''')
    doc(P, '18_Nachtrag_N01_750g_2026-09-17.docx', 'Helix Prozessautomation GmbH | Jonas Rehm', 'Angebotsnachtrag für das Format 750 g', 'N01 zu HP-2618 | 17. September 2026 | noch nicht beauftragt', '''1 Geänderte Anforderung

Auf Ihren Wunsch vom 2. September bieten wir einen zweiten Formatsatz für 750-g-Becher mit 125 mm Höhe und 95 mm Deckeldurchmesser an. Die Höhenverstellung erhält einen zusätzlichen Anschlag. Dosierdüse, Führungen und Kamerarezept werden angepasst. Für dieses Format bieten wir 72 Gutbecher je Minute am T4 an. Der Preis des Hauptpakets für 500 g bleibt unverändert.

2 Mehrpreis

Mechanische Wechselteile kosten 11.600,00 EUR netto, die zusätzliche Dosierparametrierung 4.800,00 EUR netto, Kamerarezept und Versuchsreihe 3.400,00 EUR netto sowie Einweisung und Dokumentation 4.200,00 EUR netto. Der Mehrpreis beträgt 24.000,00 EUR netto zuzüglich 4.560,00 EUR Umsatzsteuer, insgesamt 28.560,00 EUR brutto. Bei Beauftragung steigt der Hauptpreis von 675.000,00 EUR auf 699.000,00 EUR netto. Die gesonderte Vorplanung bleibt in dieser Rechnung unverändert behandelt.

3 Termin und Leistungsnachweis

Bei Freigabe bis 25. September 2026 können die Wechselteile gemeinsam gefertigt werden. Der zusätzliche Standortversuch benötigt einen Arbeitstag nach erfolgreichem 500-g-Lauf. Dieser Tag ist im bisherigen Stillstandsfenster nicht frei reserviert. Gegebenenfalls erfolgt der Versuch nach dem Produktionsstart im August 2027. Wertau stellt dafür 100 kg Produkt und 1.000 Serienbecher bereit. Eine Freigabe nach dem 25. September erfordert eine erneute Terminabstimmung.

Der Formatwechsel vom 500-g- zum 750-g-Becher soll zwei eingewiesenen Personen in höchstens 25 Minuten gelingen, ohne Reinigung und Produktwechsel. Das Nennfüllgewicht wird mit 30 aufeinanderfolgenden Bechern geprüft. Das Messfenster beträgt 744 bis 759 g bei einem Mittelwert von 750 bis 754 g. Der Nachtrag sagt keinen 90er-Takt für dieses Format zu.

4 Zahlung und Entscheidung

Wir schlagen für diesen Nachtrag 50 Prozent bei Freigabe und 50 Prozent nach erfolgreichem Standortversuch vor. Eine Bestellung des Nachtrags ersetzt nicht die noch fehlende Unterzeichnung des Hauptvertrags. Miriam Seidel bat am 16. September darum, den Nachtrag vorerst nur als Option zu führen; eine Freigabe liegt uns nicht vor.

Jonas Rehm
Geschäftsführer''')
    csvfile(P, '19_Probe_Dosierung_2026-09-03.csv', ['Probe_ID', 'Zeitpunkt', 'Rezept', 'Temperatur_C', 'Soll_g', 'Mittel_30_Becher_g', 'Minimum_g', 'Maximum_g', 'Dosiertakte_min', 'Fehlbilder_von_80', 'Bemerkung'], [
        [f'P{i:02}', f'2026-09-03T{9+(i-1)//6:02}:{((i-1)%6)*10:02}:00+02:00', '500g', t, 500, mean, low, high, rate, fail, note]
        for i, t, mean, low, high, rate, fail, note in [
            (1, '6,2', '501,2', '496,5', '505,1', 88, 0, 'Start nach Entlüften'),
            (2, '6,4', '501,0', '495,8', '506,2', 90, 0, 'weiße Deckel'),
            (3, '6,5', '501,5', '496,2', '506,7', 90, 1, 'Reflex am Siegelrand'),
            (4, '6,7', '502,1', '497,0', '507,2', 92, 1, 'ein Becher über 507 g'),
            (5, '6,8', '501,8', '496,0', '506,8', 92, 0, 'Beleuchtung geändert'),
            (6, '7,0', '501,7', '495,7', '506,1', 90, 0, 'stabil'),
            (7, '7,1', '502,2', '497,1', '506,9', 90, 0, 'neues Gebinde'),
            (8, '7,3', '502,4', '496,4', '507,4', 92, 2, 'Lufteintrag beobachtet'),
            (9, '7,5', '501,9', '496,0', '506,8', 90, 0, 'entlüftet'),
            (10, '7,6', '501,6', '495,3', '506,2', 90, 0, 'stabil'),
            (11, '7,8', '501,8', '496,6', '506,7', 90, 1, 'Etikettenmuster schräg'),
            (12, '7,9', '502,0', '496,7', '506,9', 90, 0, 'Ende vor Reinigung')]])
    doc(P, '20_Probenotiz_2026-09-03.docx', 'Helix Prozessautomation GmbH | Paula König', 'Versuchsnotiz Dosierung und Kamera', 'PV-0903 | 3. September 2026 | Werk Gersthofen', '''1 Aufbau

Anwesend waren Nora Ott für Wertau sowie Paula König und Daniel Kruse für Helix. Verwendet wurden der Vorführdosierer HP-D2, ein kurzes Versuchsband und die Optivis-Leihkamera. Der Versiegler und der Etikettierer von Wertau standen nicht im Werk. Wir hatten 120 kg Produkt aus zwei gekühlten Gebinden, 600 Leerbecher und weiße Deckel. Die Originaletikettenrolle E26-082 war noch nicht geliefert; zwölf bedruckte Muster wurden wiederholt aufgelegt.

2 Ablauf und Aufzeichnung

Von 9:00 bis 10:50 Uhr wurden zwölf Blöcke zu je 30 Bechern gewogen. Die übrige Zeit wurde Produkt zurückgeführt oder die Düse eingestellt. Die Spalte Dosiertakte_min enthält den am Dosierer eingestellten Solltakt für Einzelbecher, keine gemessene Ausbringung der Gesamtlinie. Die CSV enthält Mittelwert, Minimum und Maximum je Block. Die Tara aus zehn Leerbechern betrug im Mittel 18,4 g und wurde von den Wägewerten abgezogen.

Je Block wurden 80 markierte Fehlerbilder mit wiederverwendeten Bechern gezeigt. Fehlbilder_von_80 zählt markierte Becher, die im Versuch nicht als fehlerhaft erkannt wurden. Im Block P08 traten zwei solche Fälle auf. Nach Entlüftung und Rücknahme von 92 auf 90 Becher pro Minute wurden im nächsten Block keine übersehenen Fehler notiert. Der Höchstwert von 507,4 g in P08 ist in der Datei belassen.

3 Beobachtungen

Bei geknicktem Etikettenmuster reagierte die Erkennung auf den Beleuchtungswinkel. Frau Ott möchte mit der Originalrolle und nach Abnehmen und Wiederaufsetzen der Reinigungsabdeckung erneut prüfen. Diese Probe umfasste keinen Reinigungslauf, keinen Schichtwechsel und keine Ausschleusung in die spätere Box. Eine Freigabe zum Serienbetrieb wurde heute nicht erklärt.

Paula König, 3. September 2026
Nora Ott erhielt die CSV am selben Tag um 15:20 Uhr.''')
    mail(P, '21_2026-09-18_Vorplanung_Anrechnung.eml', 'Miriam Seidel <m.seidel@wertau-lebensmittel.de>', 'Jonas Rehm <j.rehm@helix-prozessautomation.de>', '2026-09-18T10:06:00', 'Erste Rate und bezahlte Vorplanung', '''Sehr geehrter Herr Rehm,

in Ihrer Zahlungsdatei stehen für die erste Rate weiterhin 135.000 EUR netto. Wir haben die Vorplanung schon mit 11.305 EUR brutto bezahlt. Ich hatte Ihre Aussage vom 5. August so verstanden, dass diese Zahlung bei Beauftragung des Hauptprojekts abgezogen wird. Im LOI steht ausdrücklich, dass die Vorplanung im Projektbudget berücksichtigt wird.

Bei dieser Behandlung wären aus der ersten Rate noch 125.500 EUR netto beziehungsweise 149.345 EUR brutto offen. Das ist unsere Rechnung, noch keine Freigabe einer Rate. Bitte schicken Sie mir Ihre Preisüberleitung. Unser Gesamtbudget beträgt 675.000 EUR netto einschließlich der Vorplanung; den 750-g-Nachtrag hat Frau Färber noch nicht genehmigt.

Freundliche Grüße
Miriam Seidel
Leitung Einkauf | Wertau Lebensmittel GmbH
Am Kieswerk 14, 86167 Augsburg''')
    mail(P, '22_2026-09-18_Antwort_Vorplanung.eml', 'Jonas Rehm <j.rehm@helix-prozessautomation.de>', 'Miriam Seidel <m.seidel@wertau-lebensmittel.de>', '2026-09-18T14:32:00', 'AW: Erste Rate und bezahlte Vorplanung', '''Sehr geehrte Frau Seidel,

ich erinnere das Gespräch anders. Die Vorplanung war gesondert vergütet, damit wir vor dem Anlagenauftrag anfangen konnten. Bei unserer Kalkulation von 675.000 EUR netto haben wir diese bereits erledigten Arbeiten nicht erneut angesetzt. Der Satz im LOI bedeutet nach meinem Verständnis daher nicht, dass nochmals 9.500 EUR vom angebotenen Preis abgehen.

Die fünf Preisanteile aus unserem Angebot ergeben zusammen 675.000 EUR. In keinem davon steckt die Vorplanungspauschale. Aus unserer Sicht kommen die bereits gezahlten 9.500 EUR netto hinzu. Ich sehe, dass das nicht zu Ihrer Budgetfreigabe passt. Lassen Sie uns das mit Frau Färber in der nächsten Runde besprechen; ich möchte die Frage nicht durch eine Rechnung vorwegnehmen.

Bis dahin bleibt die erste Rate nur ein Vorschlag. Es gibt weder einen unterschriebenen Hauptvertrag noch eine Rechnung über die Rate.

Freundliche Grüße
Jonas Rehm
Geschäftsführer | Helix Prozessautomation GmbH
An der Werkhalle 8, 86368 Gersthofen''')
    mail(P, '23_2026-09-21_Software_Projektende.eml', 'Dr. Eva Färber <e.faerber@wertau-lebensmittel.de>', 'Jonas Rehm <j.rehm@helix-prozessautomation.de>', '2026-09-21T11:48:00', 'Vor der nächsten Vertragsrunde: Weiterbetrieb und Abbruchkosten', '''Sehr geehrter Herr Rehm,

Herr Brandt hat mir erklärt, dass ein Ersatzrechner allein noch keine laufende Kamera ergibt. Wir müssen die Linie im Störungsfall auch ohne Wartezeit bis zum nächsten Vormittag wieder starten können. Bitte bieten Sie die Lizenzreserve und eine Unterstützung bis 22 Uhr an. Diese Kosten gehören für mich vor die Investitionsentscheidung, auch wenn sie später separat abgerechnet werden.

Sollte sich herausstellen, dass der Start im August 2027 nicht mehr erreichbar ist, brauchen wir einen nachvollziehbaren Weg aus dem Projekt. Ich möchte vor einer weiteren Zahlung die bereits bezahlten Zeichnungen und Baugruppen übernehmen können und einen anderen Integrator mit der Fertigstellung beauftragen dürfen. Wie schnell würden Sie uns die Daten herausgeben, und wie stellen Sie nicht mehr stornierbare Bestellungen nach?

Zum Haftungsabschnitt: Eine verdorbene Charge ist für uns kein entfernter Folgeschaden, sondern ein tägliches Betriebsrisiko. Den vorliegenden Ausschluss möchte ich so nicht unterschreiben. Bitte bringen Sie zu unserem Termin die Rückmeldung Ihres Versicherers mit.

Freundliche Grüße
Dr. Eva Färber
Geschäftsführerin | Wertau Lebensmittel GmbH
Am Kieswerk 14, 86167 Augsburg''')
    mail(P, '24_2026-09-22_Verhandlungsstand.eml', 'Jonas Rehm <j.rehm@helix-prozessautomation.de>', 'Miriam Seidel <m.seidel@wertau-lebensmittel.de>', '2026-09-22T15:16:00', 'Unterlagenstand für die Fortsetzung HP-2618', '''Guten Tag Frau Seidel,

ich habe Ihre Anmerkungen und die Nachricht von Frau Dr. Färber erhalten. Der Vertragsentwurf vom 9. September bleibt unser Arbeitsstand; eine neue abgestimmte Fassung gibt es noch nicht. Frau König prüft den zweiten Leistungslauf und die Terminfolge nach der verspäteten Stromversorgung. Die freie Nordfläche am 12. Juli ist weiterhin nicht bestätigt.

Optivis hält seine Termine bis 25. September vorläufig vor. Wir haben weder dort bestellt noch Dosierachsen ausgelöst. Das Angebot N01 für 750 g bleibt eine nicht bestellte Option. Zur Anrechnung der Vorplanung und zur Verlagerung der 30-Prozent-Stufe auf den Leistungslauf haben wir noch keine Verständigung.

Die Rückmeldung zum Versicherungsschutz liegt mir noch nicht vor. Einen rund um die Uhr verfügbaren Kameradienst kann ich derzeit nicht zusagen. Wir prüfen das Zeitfenster bis 22 Uhr und den zweiten Rechner. Bitte geben Sie den Entwurf noch nicht als Bestellgrundlage an Ihre Buchhaltung.

Freundliche Grüße
Jonas Rehm
Geschäftsführer | Helix Prozessautomation GmbH
An der Werkhalle 8, 86368 Gersthofen''')


def distribution():
    doc(V, '01_Besuchsnotiz_Bologna_2026-07-09.docx', 'Weserblick Messtechnik GmbH | Jan Petersen, Exportvertrieb', 'Besuch bei Luminara in Bologna', 'VP-0709 | 9. Juli 2026 | an Anneke Martens', '''1 Gespräch und Besichtigung

Ich war am 8. Juli von 10 bis 15 Uhr bei Luminara Strumenti S.r.l., Via del Calibratore 22, 40138 Bologna. Teilgenommen haben Giulia Bellini als Geschäftsführerin, Marco Rinaldi als Vertriebsleiter und Luca Serra aus dem Service. Der Kontakt war bei unserem Messegespräch in Parma am 27. Mai entstanden. Luminara verkauft derzeit Handmessgeräte zweier anderer Hersteller und möchte die stationären Module M24 und M48 ergänzen.

Das Lager umfasst drei Regalgänge, einen gesonderten Rückläuferbereich und einen Tisch für elektrische Funktionsprüfungen. Luca hat ein Netzgerät, zwei Referenzsensoren und ein Handmessgerät vorgeführt. Eine Klimakammer oder einen automatischen Prüfstand für unsere Kalibrierfolge gibt es dort nicht. Bei M48 müsste eine Kalibrierung daher vorerst in Bremen erfolgen.

2 Geschäftsmodell

Giulia möchte Ware auf eigene Rechnung kaufen, in Bologna lagern und an italienische Industriekunden weiterverkaufen. Die Rechnung an den Endkunden soll von Luminara kommen. Sie will die Verkaufspreise selbst festlegen und das Risiko verspäteter Kundenzahlungen tragen. Marco hat im Gespräch trotzdem mehrmals von seiner commission gesprochen. Auf Nachfrage meinte er damit die Differenz zwischen Einkauf und Verkauf, nicht eine laufende Zahlung von uns. Für von ihm betreute direkte Großkundengeschäfte erwartet er allerdings eine gesonderte Vergütung.

Als Startbestand wurden 60 M24 und 15 M48 genannt. Giulia denkt an 45 Tage Zahlungsziel ab tatsächlicher Lieferung nach Bologna. Ich habe erklärt, dass unser Erstangebot eine Anzahlung vorsieht. Eine Zusage zu Kredit oder Rücknahme des Bestands habe ich nicht gegeben.

3 Gebiet und Kunden

Luminara möchte Italien für drei Jahre exklusiv bearbeiten. Österreich, Schweiz und Deutschland waren nicht Gegenstand der Gespräche. Ich habe die bestehenden Direktgeschäfte mit Adria Cartoni S.p.A. ausdrücklich erwähnt. Deren italienische Werke bestellen über einen Rahmenvertrag in Bremen. Marco möchte keine neuen Werke oder neuen Warengruppen unter denselben Vorbehalt fallen lassen und erwartet, dass Bestellungen über unseren Webshop bei italienischer Lieferadresse an ihn gehen.

4 Service und weitere Unterlagen

Luca kann Wareneingang, Kabelprüfung und einen Austausch vor Ort übernehmen. Er kann keine verbindliche Fehlerdiagnose unserer Elektronik ohne Schulung leisten. Giulia möchte trotzdem innerhalb von 48 Stunden ein Austauschgerät anbieten und bittet um sechs zusätzliche Servicegeräte. Ich habe ein Schulungsangebot, eine getrennte Aufstellung der Direktkunden und einen deutschen Vertragsentwurf angekündigt. Luminara wird einen Monatsforecast und ihre Einkaufskalkulation schicken.

Jan Petersen
Reisekostenstelle Export Italien; Rückkehr nach Bremen am 9. Juli 2026.''')
    mail(V, '02_2026-07-13_Nach_dem_Besuch.eml', 'Jan Petersen <j.petersen@weserblick-messtechnik.de>', 'Giulia Bellini <g.bellini@luminara-strumenti.it>', '2026-07-13T09:18:00', 'Italien-Vertrieb: Angebot und nächste Unterlagen', '''Guten Tag Frau Bellini,

vielen Dank für die Zeit in Bologna. Ich habe Frau Martens über unser Gespräch informiert. Wir rechnen mit Luminara als Käuferin und Wiederverkäuferin, nicht mit Bestellungen im Namen von Weserblick. Ihre Preise an die Endkunden und Ihr Inkasso würden Sie selbst führen. Für die Direktkunden bleibt die Abgrenzung noch offen.

Wir senden Ihnen unser Erstangebot für 60 M24 und 15 M48 mit den am Besuchstag besprochenen Einkaufspreisen. Die Exklusivität für Italien und ein Start zum 1. Januar 2027 müssen wir im Vertrag festhalten; aus dem Besuch entsteht noch keine Gebietsreservierung. Frau Martens möchte Jahresmindestbezüge vereinbaren und einen Überblick über Ihre Servicekapazität erhalten.

Sie können Ihren Gegenvorschlag gern auf Englisch schicken. Ich lese Deutsch und Englisch; unser Vertragsentwurf kommt zunächst auf Deutsch. Die Präsentationsgeräte aus der Leihe im Juni gehören noch uns und sollten in Ihrer Lagerliste getrennt erscheinen.

Freundliche Grüße
Jan Petersen
Exportvertrieb | Weserblick Messtechnik GmbH
Am Speicherbogen 18, 28197 Bremen''')
    doc(V, '03_Angebot_WM-IT-260714.docx', 'Weserblick Messtechnik GmbH | Am Speicherbogen 18 | 28197 Bremen', 'Angebot für die Erstbevorratung Italien', 'WM-IT-260714 | 14. Juli 2026 | an Luminara Strumenti S.r.l., Bologna', '''1 Geräte und Preise

Wir bieten 60 Messmodule M24 zu je 420,00 EUR netto, zusammen 25.200,00 EUR, und 15 Messmodule M48 zu je 690,00 EUR netto, zusammen 10.350,00 EUR, an. Der Nettowarenwert beträgt 35.550,00 EUR. Die Geräte werden mit Anschlussstecker, deutscher und englischer Kurzanleitung sowie seriennummernbezogenem Werksprüfblatt geliefert. Eine italienische Kurzanleitung wird nach gemeinsamer Freigabe beigefügt.

M24 ist ein einkanaliges Modul mit 24-V-Versorgung; M48 hat zwei getrennte Messkanäle und erweiterten Ereignisspeicher. Die vereinbarten Preise gelten für die angebotenen Mengen und Lieferungen im ersten Vertragsjahr 2027. Sie sind Einkaufspreise für Luminara, keine Vorgabe für deren Wiederverkaufspreise.

2 Lieferung und Zahlung

Wir schlagen 30 Prozent Anzahlung nach Auftragsbestätigung, 40 Prozent vor Versand und 30 Prozent innerhalb von 15 Kalendertagen nach Lieferung vor. Auf den angebotenen Warenwert entfallen damit 10.665,00 EUR, 14.220,00 EUR und 10.665,00 EUR. Alle Beträge verstehen sich vor einer gegebenenfalls noch zu ergänzenden Steuerposition. Das gewünschte Zahlungsziel von 45 Tagen ab Lieferung ist nicht Bestandteil dieses Angebots.

Die Lieferung soll nach bestätigter Bestellung und Eingang der Anzahlung innerhalb von sechs Wochen erfolgen. Für einen Vertriebsstart am 1. Januar 2027 sehen wir eine Anlieferung in Bologna am 11. Januar 2027 vor, sofern die Bestellung bis 20. November 2026 bestätigt wird. Verpackung und ein Standardtransport nach Bologna sind im Warenpreis enthalten. Die Ware wird transportversichert versandt; Luminara stellt die Entladung am Lieferort.

3 Voraussetzungen

Die Darstellung mit einer Steuerposition von 0 Prozent steht unter dem Vorbehalt, dass vor endgültiger Rechnungsstellung die erforderlichen Unternehmens- und Liefernachweise vorliegen und unsere Buchhaltung die Rechnungsdaten freigibt. Eine steuerliche Kennung wird in diesem Angebot nicht angegeben. Die endgültige Rechnung wird erst nach Klärung der Stammdaten erstellt.

4 Vertriebsbeziehung

Eine Exklusivität folgt aus diesem Einzelangebot nicht. Die Parteien verhandeln gesondert über einen auf drei Jahre angelegten Vertrieb in Italien, Mindestbezüge, Kundendienst, Onlinebestellungen und die Behandlung vorhandener Direktkunden. Das Angebot bleibt bis 30. September 2026 offen. Die Präsentationsgeräte aus dem Leihvorgang WM-L-2606 sind weder in den Mengen noch im Warenwert enthalten.

Jan Petersen
Exportvertrieb''')
    doc(V, '04_Herstellervertrag_DE_2026-08-20.docx', 'Weserblick Messtechnik GmbH', 'Vertriebsvertrag Italien', 'WM-IT V1 | 20. August 2026 | Herstellerentwurf, nicht unterzeichnet', '''1 Vertragsparteien und Vertragsprodukte

1.1 Parteien

Die Weserblick Messtechnik GmbH, Am Speicherbogen 18, 28197 Bremen, Deutschland, vertreten durch Geschäftsführerin Anneke Martens, wird nachfolgend Herstellerin genannt. Die Luminara Strumenti S.r.l., Via del Calibratore 22, 40138 Bologna, Italien, vertreten durch Geschäftsführerin Giulia Bellini, wird nachfolgend Händlerin genannt. Beide Parteien führen ihr Geschäft rechtlich und wirtschaftlich selbstständig.

1.2 Produkte

Gegenstand sind die Messmodule M24 und M48 samt den im jeweiligen Angebot bezeichneten Anschlusssteckern und Prüfblättern. Andere Gerätereihen und zukünftige Nachfolgeprodukte werden nur durch schriftliche Ergänzung einbezogen. Änderungen an Gehäuse, Messbereich oder Schnittstellen kündigt die Herstellerin sechs Monate vor der ersten geänderten Lieferung an, soweit nicht eine dringende technische Korrektur erforderlich ist.

2 Einkauf und Wiederverkauf

2.1 Eigenes Geschäft

Die Händlerin kauft Produkte im eigenen Namen und auf eigene Rechnung und verkauft sie in eigenem Namen an gewerbliche Kunden weiter. Sie schließt keine Verträge im Namen der Herstellerin und darf diese weder zu Lieferungen noch zu Preisnachlässen verpflichten. Kundenforderungen und das Risiko ausbleibender Kundenzahlungen verbleiben bei der Händlerin. Sie richtet ihre Wiederverkaufspreise selbst ein.

2.2 Vergütung

Die Händlerin erzielt ihre Vergütung aus der Differenz zwischen ihrem Einkaufspreis und ihrem Wiederverkaufserlös. Für eigene Wiederverkäufe zahlt die Herstellerin keine Provision. Wünscht die Herstellerin Leistungen der Händlerin für ein unmittelbares Kundengeschäft, werden Umfang und Entgelt vor Beginn in einem gesonderten Serviceauftrag festgehalten. Die bloße Weitergabe einer Anfrage begründet in diesem Entwurf keinen Zahlungsanspruch.

3 Gebiet und Direktkunden

3.1 Italien

Die Herstellerin bestellt die Händlerin für die Dauer dieses Vertrags exklusiv für das Gebiet Italien. Deutschland, Österreich und die Schweiz gehören nicht zum Vertragsgebiet. Die Herstellerin wird in Italien keinen weiteren Händler für die Vertragsprodukte bestellen. Die Händlerin unterhält in Bologna einen erreichbaren Ansprechpartner und ein Lager für die laufende Versorgung italienischer Kunden.

3.2 Vorbehaltene Kunden

Die Herstellerin behält sich die direkte Belieferung von Adria Cartoni S.p.A. und sämtlichen gegenwärtigen und zukünftigen Konzerngesellschaften dieses Unternehmens vor. Der Vorbehalt erstreckt sich auf alle Standorte und alle Vertragsprodukte. Anfragen dieser Unternehmen leitet die Händlerin unverzüglich an die Herstellerin weiter. Die Anlage DK-IT-01 enthält die derzeit bekannten italienischen Abnahmestellen; sie begrenzt den Vorbehalt nicht auf diese Stellen.

3.3 Weitere Direktgeschäfte

Die Herstellerin darf ferner europaweite Rahmenverträge mit zentralem Einkauf außerhalb Italiens auch für italienische Lieferorte erfüllen. Sie informiert die Händlerin vor dem ersten Abruf, soweit die Vertraulichkeit des jeweiligen Rahmenvertrags dies zulässt. Eine Anrechnung dieser Umsätze auf den Mindestbezug der Händlerin oder eine Vergütung ist nicht vorgesehen. Italienische Einzelanfragen außerhalb dieser Vorbehalte werden der Händlerin zur Bearbeitung zugeleitet.

---

4 Onlinevertrieb und Außenauftritt

4.1 Bestellwege

Die Händlerin darf einen eigenen italienischsprachigen Onlineshop betreiben. Die Nutzung offener Marktplätze für Vertragsprodukte setzt eine vorherige schriftliche Zustimmung der Herstellerin voraus. Die Händlerin soll in ihrem Shop Bestellungen mit Lieferadressen außerhalb Italiens sperren und Anfragen von außerhalb des Gebiets an die Herstellerin weitergeben. Die Herstellerin entscheidet über deren Bearbeitung. Diese Vorgaben gelten auch für bezahlte Onlinewerbung, die gezielt auf andere Länder ausgerichtet wird.

4.2 Herstellerportal

Das Bestellportal der Herstellerin bleibt für bestehende Benutzer geöffnet. Bestellungen italienischer Lieferadressen außerhalb der Direktkundenvorbehalte werden intern an die Händlerin weitergeleitet, sobald die erforderliche Zuordnung im Portal eingerichtet ist. Bis dahin bestätigt die Herstellerin solche Aufträge selbst. Eine technische Umstellung wird nicht vor dem 1. April 2027 zugesagt. Umsätze dieser Übergangsphase werden nicht auf den Mindestbezug angerechnet.

4.3 Kennzeichen und Werbung

Die Händlerin darf Namen und Produktabbildungen der Herstellerin für den Absatz der Vertragsprodukte verwenden. Sie kennzeichnet ihren eigenen Firmennamen und ihre eigene Verantwortung für den Verkauf. Technische Werbeaussagen müssen mit den überlassenen Datenblättern übereinstimmen. Übersetzungen sicherheitsrelevanter Hinweise werden der Herstellerin vor Veröffentlichung zur Freigabe vorgelegt; die Herstellerin antwortet innerhalb von zehn Arbeitstagen.

5 Mindestbezug und Planung

5.1 Jahresbeträge

Die Händlerin bezieht im Kalenderjahr 2027 Vertragsprodukte mit einem Nettowarenwert von mindestens 180.000,00 EUR, im Jahr 2028 mindestens 220.000,00 EUR und im Jahr 2029 mindestens 260.000,00 EUR. Maßgeblich sind ausgelieferte und nicht gutgeschriebene Warenlieferungen an die Händlerin. Serviceentgelte, Steuer, Frachtzuschläge und Direktkundengeschäfte zählen nicht mit. Die Erstbevorratung von 35.550,00 EUR zählt zum Mindestbezug 2027, wenn sie 2027 ausgeliefert wird.

5.2 Folge einer Unterschreitung

Wird der Jahresmindestbezug um mehr als zehn Prozent unterschritten, darf die Herstellerin die Exklusivität für das Folgejahr durch Mitteilung bis 31. Januar dieses Folgejahres beenden. Zuvor besprechen die Parteien die Bezugszahlen und Lieferausfälle der Herstellerin. Nicht rechtzeitig gelieferte, bestätigte Bestellungen werden bei dieser Berechnung so behandelt, als seien sie im zugesagten Jahr ausgeliefert worden. Eine Pflicht zur Bezahlung nicht bestellter Ware entsteht durch den Mindestbezug allein nicht.

5.3 Forecast

Die Händlerin übermittelt bis zum zehnten Arbeitstag jedes Monats eine rollierende Mengenplanung für die nächsten zwölf Monate. Die ersten zwei Monate werden durch Einzelbestellungen konkretisiert; die übrigen Monate bleiben Planung. Eine Forecast-Zeile ist keine Bestellung. Lieferkapazitäten werden erst mit der jeweiligen Auftragsbestätigung verbindlich reserviert.

6 Bestellungen und Änderungen

Bestellungen nennen Produkt, Menge, gewünschten Liefertermin und Lieferort. Ein Einzelauftrag kommt durch die Auftragsbestätigung der Herstellerin zustande. Abweichungen von Menge, Preis oder Termin müssen darin ausdrücklich bezeichnet werden und bedürfen der Bestätigung durch die Händlerin. Nach bestätigtem Produktionsbeginn kann die Händlerin Mengen nicht einseitig stornieren. Die Herstellerin legt auf Anfrage die Kosten einer möglichen Änderung vor.

---

7 Preise und Zahlung

7.1 Einkaufspreise

Für Lieferungen 2027 gelten 420,00 EUR netto je M24 und 690,00 EUR netto je M48. Die Preise enthalten Standardverpackung und Standardtransport nach Bologna bei einem Nettowarenwert von mindestens 10.000,00 EUR je Sendung. Kleinere Sendungen werden mit 180,00 EUR netto Fracht berechnet. Die Preisänderung für ein Folgejahr wird bis 30. September des Vorjahres angekündigt; bestätigte Einzelaufträge bleiben unverändert.

7.2 Zahlungsstufen

Die Händlerin zahlt 30 Prozent nach Auftragsbestätigung, 40 Prozent vor Versand und 30 Prozent binnen 15 Kalendertagen nach Ablieferung. Eine abweichende Kreditlinie oder ein Ziel von 45 Tagen entsteht nicht durch das Einreichen einer Bestellung mit diesen Angaben. Steuerpositionen werden erst nach Prüfung der erforderlichen Stammdaten und Liefernachweise in der endgültigen Rechnung ausgewiesen. Die Proforma ersetzt keine Rechnung und löst für sich keine Zahlungspflicht aus.

7.3 Überfällige Beträge

Bei unbeglichenen fälligen Forderungen darf die Herstellerin nach schriftlicher Erinnerung und einer Nachfrist von zehn Kalendertagen weitere Auslieferungen aussetzen. Sie benennt die betroffenen Rechnungen und die zur Wiederaufnahme erforderliche Zahlung. Bereits bestätigte Liefertermine werden anschließend neu abgestimmt. Unstreitige Teile einer Rechnung bleiben zu zahlen, auch wenn über einen anderen Teil Einwendungen bestehen.

8 Transport, Wareneingang und Eigentum

Die Herstellerin verpackt die Geräte transportsicher und übergibt sie dem ausgewählten Frachtführer. Die Gefahr geht mit Ablieferung an die Händlerin in Bologna über; die Händlerin organisiert und verantwortet das Entladen. Sichtbare Transportschäden werden bei Empfang auf dem Ablieferbeleg festgehalten und mit Fotos binnen zwei Arbeitstagen übermittelt. Nicht sichtbare Funktionsabweichungen werden nach Feststellung mit Seriennummer und Befund gemeldet.

Die Ware bleibt bis zur vollständigen Zahlung des jeweiligen Kaufpreises Eigentum der Herstellerin. Der ordentliche Weiterverkauf durch die Händlerin ist gestattet. Vor vollständiger Bezahlung darf sie die Ware nicht als Sicherheit an Dritte übertragen. Bei Rückholung noch nicht verkaufter Ware stimmen die Parteien Zutritt, Zählung und Transport ab; die Herstellerin darf Lagerflächen der Händlerin nicht ohne Abstimmung betreten.

9 Kundendienst und Mängelbearbeitung

9.1 Erstkontakt

Die Händlerin nimmt Kundenmeldungen auf Italienisch entgegen, prüft Seriennummer, Spannungsversorgung und Anschlussbelegung und dokumentiert das Ergebnis im RMA-Protokoll. Eingriffe in vergossene Elektronik und Kalibrieränderungen sind ohne Freigabe der Herstellerin untersagt. Kann die Händlerin einen Fehler nicht eingrenzen, übermittelt sie Messwerte und Fotos an den Service in Bremen. Die Herstellerin meldet sich an Arbeitstagen innerhalb von zwei Arbeitstagen.

9.2 Reparatur

Die Herstellerin bearbeitet produktbezogene Mängel innerhalb von 24 Monaten ab Lieferung an die Händlerin durch Reparatur oder Ersatz. Die Händlerin meldet Fehler möglichst vor Rücksendung an und erhält eine RMA-Nummer. Bestätigt sich ein Fehler der gelieferten Ware, übernimmt die Herstellerin angemessene Rückfracht und erneuten Versand. Schäden durch falschen Anschluss, Flüssigkeitseintritt oder nicht freigegebene Eingriffe werden nach Befund gesondert angeboten. Eine Diagnosepauschale wird nur nach vorheriger Zustimmung der Händlerin berechnet.

---

9.3 Servicevergütung und Austauschbestand

Die Händlerin erbringt die Aufnahme und einfache Erstprüfung gegenüber eigenen Kunden aus ihrer Handelsspanne. Für Einsätze bei vorbehaltenen Direktkunden braucht sie vorab einen schriftlichen Serviceauftrag. Ein Pool aus sechs Austauschgeräten wird nicht unentgeltlich gestellt; dessen Kauf oder Leihe wird gesondert vereinbart. Die Herstellerin sagt keine Wiederherstellung innerhalb von 48 Stunden zu. Die Händlerin darf eigene kürzere Servicezusagen machen, ohne die Herstellerin damit zu verpflichten.

9.4 Schulung

Die Herstellerin schult zwei Mitarbeiter der Händlerin an zwei Tagen in Bremen zu Anschlussprüfung, Fehleraufnahme und Rücksendung. Die erste Schulung ist im Vertriebsaufbau enthalten, Reise und Unterkunft trägt die Händlerin. Eine Befugnis zur Kalibrierung entsteht erst nach gesonderter Einweisung und schriftlicher Freigabe. Technische Mitteilungen über geänderte Prüfverfahren werden dem benannten Serviceleiter übersandt.

10 Kunden- und Verkaufsdaten

Die Händlerin übermittelt monatlich eine Liste ihrer Verkäufe mit Kundenfirma, Lieferanschrift, Ansprechpartner, geschäftlicher E-Mail-Adresse, Produkt, Seriennummer, Menge, Nettoverkaufspreis und Einsatzbereich. Die Daten sollen der Betreuung, Ersatzteilversorgung und Mengenplanung dienen. Die Herstellerin führt die Liste in ihrem Kundenverwaltungssystem. Ein Zugriff des Herstellers auf offene Angebote und noch nicht abgeschlossene Verhandlungen ist in dieser Fassung ebenfalls vorgesehen. Einzelheiten enthält der Entwurf KD-IT-01.

Die Händlerin erläutert ihren Kunden die Weitergabe im Rahmen ihrer Kundenkommunikation. Nach Vertragsende erhält die Herstellerin die letzte vollständige Liste für die fortlaufende Gerätebetreuung. Eine gesonderte Vergütung für den Datenbestand ist nicht vorgesehen. Vertrauliche Informationen der jeweils anderen Partei werden nur Beschäftigten zugänglich gemacht, die sie für Einkauf, Vertrieb oder Service benötigen.

11 Dauer und vorzeitige Beendigung

11.1 Laufzeit

Der Vertrag beginnt am 1. Januar 2027 und endet am 31. Dezember 2029, ohne dass es einer Kündigung bedarf. Eine Verlängerung muss bis 30. September 2029 schriftlich vereinbart werden. Die Parteien besprechen im Juni 2029 Bezugszahlen, Serviceaufwand und eine mögliche Fortsetzung. Eine ordentliche Kündigung während der festen Laufzeit ist in diesem Entwurf nicht vorgesehen.

11.2 Pflichtverletzung

Bei wesentlicher Pflichtverletzung kann die andere Partei nach schriftlicher Benennung des Verstoßes und fruchtlosem Ablauf einer Abhilfefrist von 30 Kalendertagen vorzeitig beenden. Für ausstehende Zahlungen gilt zunächst die Nachfrist aus Ziffer 7.3; bleibt sie erfolglos, kann eine weitere Beendigungsandrohung mit zehn Kalendertagen Frist folgen. Sofortige Beendigung bleibt schweren Fällen vorbehalten, in denen eine Fortsetzung bis zum Fristablauf nicht zumutbar ist.

11.3 Auswirkungen

Bereits bestätigte Einzelbestellungen werden ausgeführt, sofern sie nicht gesondert aufgehoben werden. Offene Forderungen bleiben bestehen. Die Händlerin stellt sich nach Vertragsende nicht mehr als exklusiver Vertriebspartner dar und entfernt entsprechende Angaben binnen zehn Arbeitstagen. Technische Unterlagen für bereits ausgelieferte Produkte darf sie zur Betreuung ihrer Kunden weiterverwenden.

---

12 Restbestand und Leihgeräte

12.1 Lagerware

Die Herstellerin ist nach Vertragsende nicht verpflichtet, verkaufsfähige Lagerware zurückzukaufen. Die Händlerin darf den bei Vertragsende vorhandenen Bestand während sechs Monaten unter Verwendung zutreffender Produktbezeichnungen abverkaufen. Sie übermittelt binnen zehn Arbeitstagen eine Liste mit Seriennummern, Anschaffungsdatum und Einkaufspreis. Nach Ablauf der sechs Monate ist eine weitere Nutzung der Kennzeichen in einem Onlineshop nur nach Abstimmung mit der Herstellerin zulässig.

12.2 Freiwillige Rücknahme

Die Herstellerin kann im Einzelfall ungeöffnete, unveränderte Geräte aktueller Bauart zurücknehmen. Voraussetzung sind eine schriftliche Mengenfreigabe und ein vereinbarter Rücknahmepreis vor Versand. Ein Rücknahmeverlangen allein berechtigt die Händlerin nicht, Zahlungen aus anderen Lieferungen zurückzuhalten. Die Kosten der Rückfracht werden in der jeweiligen Freigabe festgelegt.

12.3 Leihgeräte

Geräte aus gesonderten Leihvorgängen bleiben Eigentum der Herstellerin. Sie werden mit Seriennummern gesondert erfasst und binnen 20 Kalendertagen nach Ende der Leihe zurückgegeben. Ein Leihgerät darf nicht als neuer Verkaufsbestand angeboten werden. Übliche Gebrauchsspuren aus vereinbarten Vorführungen werden im Rückgabeprotokoll beschrieben und nicht als Verkauf oder Erwerb des Geräts behandelt.

13 Haftung und Produkthinweise

Für vorsätzlich oder grob fahrlässig verursachte Schäden und schuldhafte Verletzungen von Leben, Körper oder Gesundheit haften die Parteien unbeschränkt. Bei leichter Fahrlässigkeit wird die Haftung für wesentliche Vertragspflichten auf vorhersehbare Schäden und auf den Nettowarenwert der Lieferungen der letzten zwölf Monate begrenzt. In den ersten zwölf Monaten tritt ein Höchstbetrag von 180.000,00 EUR an diese Stelle. Zwingende Haftung bleibt unberührt. Die Händlerin darf sicherheitsbezogene Informationen nicht verändern und informiert die Herstellerin unverzüglich über wiederholte gleichartige Geräteausfälle.

Rückruf- oder Korrekturmaßnahmen stimmen die Parteien anhand der betroffenen Seriennummern ab. Eine Kostenaufteilung erfolgt nach Ursache und Verantwortungsbeitrag. Die Händlerin bewahrt Zuordnungen ausgelieferter Seriennummern zu Kunden für zehn Jahre auf. Eine pauschale Übernahme sämtlicher Rückrufkosten durch die Händlerin wird nicht vereinbart.

14 Vertraulichkeit und Übertragung

Preise, unveröffentlichte Produktinformationen und die ausgetauschten Kundendaten werden vertraulich behandelt. Diese Pflicht gilt drei Jahre nach Vertragsende fort; technische Geheimnisse bleiben geschützt, solange sie nicht allgemein zugänglich sind. Eine Übertragung des gesamten Vertrags auf ein anderes Unternehmen bedarf der vorherigen Zustimmung der anderen Partei. Die Einschaltung eines eigenen Logistikdienstleisters ist gestattet, wenn dessen Zugriff auf Daten auf Versand und Lagerung begrenzt wird.

15 Recht, Sprache und Gerichtsstand

Die Herstellerin schlägt deutsches Recht unter Ausschluss des Übereinkommens über Verträge über den internationalen Warenkauf vor. Ausschließlicher Gerichtsstand soll Bremen sein. Maßgebliche Vertragssprache soll Deutsch sein. Diese Festlegungen sind Bestandteil des Herstellerentwurfs; eine Zustimmung der Händlerin zu Recht, Gericht oder Sprache liegt noch nicht vor. Die Parteien wollen eine endgültige Sprachfassung gemeinsam freigeben.

16 Abschluss

Der Vertrag wird erst durch Unterzeichnung beider Parteien verbindlich. Änderungen bedürfen einer von beiden Geschäftsführungen bestätigten Erklärung. Einkaufsbedingungen in einer späteren Bestellung werden nur Vertragsbestandteil, wenn die Herstellerin ihnen ausdrücklich zustimmt. Bei nicht durchführbaren Einzelregelungen besprechen die Parteien eine ihrem wirtschaftlichen Ziel entsprechende Ersatzregelung. Die übrigen Vereinbarungen bleiben bestehen.

Anneke Martens für Weserblick Messtechnik GmbH: Unterzeichnung steht aus.

Giulia Bellini für Luminara Strumenti S.r.l.: Unterzeichnung steht aus.''')
    doc(V, '05_Distributor_counterproposal_EN_2026-09-03.docx', 'Luminara Strumenti S.r.l. | Via del Calibratore 22 | 40138 Bologna', 'Italy distribution counterproposal', 'LS-WM-03 | 3 September 2026 | Giulia Bellini to Anneke Martens | unsigned', '''1 Purchase and resale

Luminara offers to purchase the M24 and M48 modules for resale in its own name and for its own account. Luminara will issue invoices to its customers, collect the receivables and bear their credit risk. Luminara will set its resale prices independently. Weserblick will invoice Luminara at the agreed purchase prices. No sales commission will be payable on goods purchased and resold by Luminara.

Our proposed purchase prices for 2027 are EUR 420 per M24 and EUR 690 per M48, net of any tax position still to be confirmed. The opening order would contain 60 M24 and 15 M48, giving a total goods value of EUR 35,550. The presentation units already held in Bologna under loan WM-L-2606 are not part of this order and will remain separately recorded.

2 Territory and reserved accounts

The territory will be Italy only. The initial term will run from 1 January 2027 until 31 December 2029. Weserblick will not appoint another distributor in Italy and will direct enquiries from Italian delivery addresses to Luminara, except for the reserved accounts expressly listed in the signed account schedule.

The reservation for Adria Cartoni S.p.A. should cover only its existing Parma and Verona plants and only orders under the current framework signed before this agreement. New plants, acquisitions, affiliates and new product families should not automatically become reserved accounts. Additional reservations will require Luminara's written agreement before the first direct quotation is issued.

When Luminara is asked to visit a reserved account, commission the equipment or provide first response, Weserblick will issue a separate service order. We propose EUR 95 per technician hour plus agreed travel expenses. When Luminara identifies and develops a new opportunity that Weserblick wishes to supply directly, the parties will agree a referral fee before the customer receives a direct offer. Luminara is not asking for a commission on purchases made for its own stock.

3 Internet sales

Luminara will operate an Italian web shop under its own name. It will be allowed to accept unsolicited orders from customers outside Italy without automatically blocking their addresses. Paid campaigns will focus on Italy during the first year. Product pages and delivery terms will clearly identify Luminara as the seller.

Weserblick's portal should route non-reserved Italian accounts to Luminara from the start date, not from April. If the change cannot be completed in time, the parties should agree how those sales affect Luminara's minimum purchase before the first portal order is accepted. We cannot fund an Italian stock while the same customers continue buying directly without any adjustment to our targets.

---

4 Minimum purchases and supply

Luminara proposes annual net purchases of EUR 150,000 in 2027, EUR 190,000 in 2028 and EUR 230,000 in 2029. The opening delivery counts towards 2027. Service fees and the loan units do not count. The monthly forecast we are preparing will be a planning estimate and will not turn all twelve months into binding purchase orders.

If Weserblick cannot deliver an accepted order in the agreed year, the order value will count towards that year's minimum for the purpose of exclusivity. If direct sales to non-reserved Italian customers continue, their net goods value should also count towards the target. Luminara does not propose to pay for quantities that were neither ordered nor delivered.

Weserblick should review any shortfall with Luminara before changing exclusivity. Luminara proposes a three-month recovery period, measured against an agreed order schedule. There should be no automatic loss of exclusivity based on a forecast revision or an unpaid end-customer invoice.

5 Payment and first shipment

Luminara proposes payment 45 calendar days after actual delivery in Bologna, without advance payment. For a delivery on 11 January 2027, the due date would therefore be 25 February 2027. The first order is subject to agreement of this term and cannot be confirmed under the advance-payment terms in offer WM-IT-260714.

The credit exposure can be limited to EUR 60,000 of unpaid goods invoices. If accepting another shipment would exceed that amount, the parties will agree a split shipment or an interim payment. This credit limit is an offer by Luminara, not an existing facility. Goods remain covered by the agreed retention of title, but Luminara must remain free to resell them in the ordinary course of business.

6 Warranty and service

Luminara will provide Italian-language intake, check the supply voltage and cabling and record the serial number and observed fault. It will not open sealed electronics or perform calibration without training and written authorisation. For confirmed product defects, Weserblick should provide repair or replacement and cover reasonable return freight. Luminara proposes that the 24-month period start with its sale to the end customer, but no later than six months after delivery to Bologna.

We need a separately funded exchange pool if we are to offer a replacement within 48 hours. We propose four M24 and two M48 on loan, with ownership remaining with Weserblick. Work beyond intake and a basic connection check should be charged at the agreed hourly rate when Weserblick requests it. The distributor margin should not silently include repeated site visits for a manufacturing fault. Weserblick should not be bound by any separate paid service promise made by Luminara to its own customer.

---

7 Customer information

Luminara will provide monthly product quantities, serial numbers, end-customer company names, delivery towns and application sectors for delivered devices. Personal contact information will be provided where needed for an identified service case or with a customer introduction agreed by Luminara. We do not agree to a standing transfer of all named contacts, private mobile numbers, our resale prices or open opportunities.

Weserblick should use the transmitted information for product support, traceability and production planning. It should not use the list to quote directly to Luminara's active accounts during the term. The parties need an agreed retention period and access list before uploading any customer file. The current data-field proposal KD-IT-01 is not approved by Luminara.

8 Stock at the end of the relationship

If the agreement expires after three years or Weserblick ends it without a material breach by Luminara, Weserblick should repurchase unopened, current-model units delivered in the preceding twelve months at the original net purchase price. Luminara will provide serial numbers and purchase invoices within 15 days. Weserblick will arrange collection within 30 days and pay within 30 days of receipt. Units with damaged packaging can be reviewed individually; unopened accessories should follow the same process.

Older saleable units should have a twelve-month sell-off period. Luminara must be allowed to identify the products truthfully in its web shop during that period. We do not propose a repurchase at the normal price for worn presentation units, customer returns or devices with unapproved modifications. Loan devices will be returned under their separate loan terms and must not be counted twice in any stock settlement.

9 Termination and transition

Either party may terminate for a material breach that remains unremedied 30 days after a written notice describing the breach. Before suspending deliveries for non-payment, Weserblick should identify the overdue invoice and allow the agreed cure period. A dispute about a specific RMA should not suspend unrelated fully paid orders.

During the final three months, the parties should identify open customer orders, warranty cases and stock serial numbers. Accepted purchase orders will continue unless both parties agree otherwise. Luminara will stop presenting itself as the exclusive distributor when the term ends, while retaining product documentation needed to support already sold devices.

10 Applicable law and forum

Luminara proposes Italian law and the courts of Bologna. We have not accepted German law, Bremen jurisdiction or German as the only authoritative contract language. We are willing to discuss a bilingual German and English execution version, with a jointly agreed rule for any discrepancy. No choice of law, court or controlling language has yet been agreed.

Giulia Bellini
Managing Director, Luminara Strumenti S.r.l.
This counterproposal remains open for discussion and has not been signed by either party.''')
    mail(V, '06_2026-09-04_Counterproposal_cover.eml', 'Giulia Bellini <g.bellini@luminara-strumenti.it>', 'Jan Petersen <j.petersen@weserblick-messtechnik.de>', '2026-09-04T10:21:00', 'Our counterproposal and the opening stock', '''Dear Jan,

Please find our counterproposal dated 3 September. The attachment covers the commercial terms we discussed, including the distinction between our own resale margin and paid work on your direct accounts. Marco used the word commission at the visit, but he was referring to margin on the devices we buy. We will invoice our customers ourselves and carry their payment risk.

The EUR 35,550 opening stock is feasible only if payment is due 45 days after delivery. Our customers normally pay after installation, so paying 70 percent before the goods leave Bremen would use most of the money reserved for the launch. We can discuss a EUR 60,000 credit limit, but no facility has been approved on either side.

I also need a closed list of reserved accounts. The phrase future group companies is too broad for our sales plan. Please keep the opening order on hold while we resolve these points; the quantity request is not acceptance of your July payment terms.

Best regards,
Giulia Bellini
Managing Director | Luminara Strumenti S.r.l.
Via del Calibratore 22, 40138 Bologna, Italy''')
    mail(V, '07_2026-09-07_Intern_Direktkunden.eml', 'Anneke Martens <a.martens@weserblick-messtechnik.de>', 'Jan Petersen <j.petersen@weserblick-messtechnik.de>', '2026-09-07T08:44:00', 'Luminara: Adria und Mindestbezug', '''Guten Morgen Jan,

Adria Cartoni bleibt bei uns. Der Rahmenvertrag läuft über Bremen, und im letzten Jahr kamen 112.000 EUR Nettoumsatz aus Parma und Verona. Der Einkauf hat einen zusätzlichen Standort bei Bari angekündigt, aber noch nichts bestellt. Ich möchte den Konzern nicht in zwei Betreuungsmodelle teilen. Bitte kennzeichne Bari in der Liste als Gespräch, nicht als bestehenden Abnehmer.

Für die Exklusivität hatte ich 180.000 EUR im ersten Jahr angesetzt. Luminara schlägt nur 150.000 vor. Sobald der Forecast da ist, möchte ich sehen, ob die Startlieferung darin schon steckt. Sie darf nicht zweimal zum Mindestbezug addiert werden. Direktlieferungen an Adria sollen nach unserem Entwurf weiterhin nicht mitzählen.

Mit Provisionen auf ihre eigenen Einkäufe fangen wir nicht an. Wenn Luca zu einem unserer Kunden fährt, können wir einen Serviceauftrag erteilen. Für bloße Namen in einer Kontaktliste habe ich noch keine Vergütung zugesagt. Bitte keine endgültige Auftragsbestätigung aus dem Portal schicken, solange Zahlung und Exklusivität ungeklärt sind.

Viele Grüße
Anneke Martens
Geschäftsführerin | Weserblick Messtechnik GmbH
Am Speicherbogen 18, 28197 Bremen''')
    doc(V, '08_Direktkundenanlage_DK-IT-01.docx', 'Weserblick Messtechnik GmbH | Jan Petersen', 'Direktkunden und italienische Lieferstellen', 'DK-IT-01 V1 | 8. September 2026 | Anlage zum Herstellerentwurf', '''1 Adria Cartoni S.p.A.

Die Kundennummer AC-100 führt den zentralen Einkauf von Adria Cartoni S.p.A. mit unmittelbarer Abrechnung durch Weserblick. Die seit 2024 bedienten Lieferstellen sind Werk Parma, interne Standortnummer AC-PAR, und Werk Verona, interne Standortnummer AC-VER. Geliefert werden M24 und M48. Der Nettowarenumsatz dieser beiden Lieferstellen betrug 2025 zusammen 112.000,00 EUR. Für 2027 liegt noch kein verbindlicher Jahresabruf vor.

Für einen Standort bei Bari, Kennung im Vertrieb AC-BAR-P, wurde am 19. August 2026 eine technische Auskunft angefordert. Es gibt dazu weder einen Auftrag noch eine freigegebene Lieferadresse. Die Aufnahme in diese Gesprächsliste bezeichnet daher keine erfolgte Belieferung. Anneke Martens möchte auch zukünftige Adria-Standorte unter den Direktkundenvorbehalt fassen; Luminara hat dem nicht zugestimmt.

2 Andere italienische Kunden

Die Unternehmen Officine Roveta S.r.l. in Modena und Vetraria Selce S.r.l. in Vicenza haben 2026 Geräte über das Herstellerportal angefragt. Sie besitzen ein Benutzerkonto, sind aber nicht als vorbehaltene Direktkunden aufgeführt. Officine Roveta fragt derzeit acht M24 an, Vetraria Selce vier M48. Die Anfragen sind nicht bestätigt. Jan Petersen hat beide Kontakte am 8. September in den noch offenen Portalumstellungsplan aufgenommen.

3 Abgrenzung im Gespräch

Giulia Bellini verlangt eine abschließende Aufzählung einzelner Standorte und eine Begrenzung auf den vorhandenen Adria-Rahmenvertrag. Weserblick möchte den Konzernvorbehalt und zusätzlich europaweite Rahmenverträge mit zentralem Einkauf außerhalb Italiens beibehalten. Eine gemeinsame Anlage wurde bislang nicht unterschrieben. Diese Liste gibt den von Weserblick am 8. September versandten Stand wieder.

Jan Petersen
Exportvertrieb''')
    doc(V, '09_Produkt_und_Serviceblatt_2026-08-24.docx', 'Weserblick Messtechnik GmbH | Ralf Döring, Serviceleitung', 'Gerätedaten und Serviceumfang M24 und M48', 'TB-IT-0824 | 24. August 2026 | an Luca Serra, Luminara', '''1 Geräteausführung

M24 verarbeitet einen Messkanal mit Eingang 4 bis 20 mA bei 24 V Gleichspannung. Das Gehäuse misst 118 mal 76 mal 42 mm. Das Gerät wird auf einer Tragschiene befestigt. M48 verarbeitet zwei getrennte Kanäle, besitzt einen lokalen Ereignisspeicher für 10.000 Datensätze und misst 148 mal 86 mal 48 mm. Beide Geräte sind für geschützte Schaltschränke bei 0 bis 45 Grad Celsius vorgesehen. Sie sind weder für das Abspritzen noch für den unmittelbaren Einbau in Nassbereiche angeboten.

Die Lieferung enthält einen codierten Anschlussstecker, Kurzanleitung und Werksprüfblatt mit Seriennummer. Der dokumentierte maximale Anzeigefehler beträgt bei 23 Grad Celsius 0,2 Prozent des Endwerts. Leitungen, Sensoren und eine Integration in die Steuerung des Endkunden sind nicht im Gerätepreis enthalten. Das Gerät darf nicht für eine sicherheitsgerichtete Abschaltung eingesetzt werden.

2 Wareneingang und Erstprüfung

Luminara prüft Verpackung, Seriennummer und Zubehör unmittelbar bei Eingang. Bei einer Fehlermeldung werden Versorgungsspannung ohne und mit angeschlossenem Gerät sowie der Eingangsstrom dokumentiert. Das Gehäuse bleibt geschlossen. Vor einem Austausch werden Parametrierdatei und Fehlerspeicher gesichert, soweit das Gerät noch erreichbar ist. Eine Kalibrieränderung oder Reparatur an der Leiterplatte darf nur nach gesonderter Freigabe erfolgen.

3 Bearbeitung in Bremen

Der Service vergibt eine RMA-Nummer und bittet um Fehlerbeschreibung, Betriebsdauer, Anschlussfoto und Herkunft des verwendeten Netzteils. Die Wareneingangsprüfung in Bremen benötigt normalerweise fünf Arbeitstage ab Zugang; ein Reparaturtermin wird nach Befund genannt. Für bestätigte Produktmängel wird keine Diagnose berechnet. Bei Fremdursachen erhält Luminara ein Angebot vor Reparatur. Die Pauschale für eine zusätzlich beauftragte Diagnose beträgt 85,00 EUR netto je Gerät.

4 Austausch und Vor-Ort-Arbeit

Ein Austausch innerhalb von 48 Stunden kann nur aus einem örtlichen Gerätebestand erfolgen. Derzeit sind vier M24 und zwei M48 als Präsentationsleihgeräte in Bologna; diese sind nicht als frei verkäuflicher Austauschbestand freigegeben. Ein Wechsel von Präsentation zu Servicepool setzt eine neue Zuordnung der Seriennummern voraus. Weserblick bietet bislang keinen unentgeltlichen Austauschpool an.

Für Service bei einem Direktkunden vereinbart Weserblick vor dem Besuch einen Stundensatz und eine Kostenobergrenze. Luminara hat 95,00 EUR netto je Stunde vorgeschlagen; dieser Satz ist noch nicht bestätigt. Reisezeiten und Material sollen im jeweiligen Auftrag gesondert stehen. Eigene Serviceversprechen gegenüber Luminara-Kunden sind nicht Bestandteil unserer bisherigen Gerätekalkulation.

Ralf Döring
Serviceleitung | r.doering@weserblick-messtechnik.de''')
    doc(V, '10_Proforma_PF-IT-260909.docx', 'Weserblick Messtechnik GmbH | Am Speicherbogen 18 | 28197 Bremen', 'Proforma für die Verhandlung der Erstlieferung', 'PF-IT-260909 | 9. September 2026 | keine Zahlungsanforderung', '''Empfängerin und vorgesehene Lieferstelle: Luminara Strumenti S.r.l., Via del Calibratore 22, 40138 Bologna, Italien. Ansprechpartnerin: Giulia Bellini. Bezug: Angebot WM-IT-260714 und Mengenabstimmung beim Besuch am 8. Juli 2026.

1 Warenpositionen

Position 1: Messmodul M24, 60 Stück zu je 420,00 EUR netto. Positionswert: 25.200,00 EUR.

Position 2: Messmodul M48, 15 Stück zu je 690,00 EUR netto. Positionswert: 10.350,00 EUR.

Nettowarenwert: 35.550,00 EUR. Standardverpackung und Standardtransport nach Bologna: im Warenpreis enthalten. Steueransatz für diese Proforma: 0 Prozent, 0,00 EUR. Proforma-Gesamtbetrag: 35.550,00 EUR.

2 Nachweisvorbehalt

Der angesetzte Wert von 0 Prozent gilt hier ausschließlich als Verhandlungsannahme und steht unter Nachweisvorbehalt. Die erforderlichen Unternehmens-, Abnehmer- und Liefernachweise sind noch nicht vollständig abgeglichen. Eine steuerliche Kennung wird nicht angegeben. Vor endgültiger Rechnungsstellung prüft die Buchhaltung die Nachweise und gibt die Rechnungsdaten frei. Mit dieser Proforma wird keine abschließende steuerliche Behandlung bestätigt.

3 Zahlung und Lieferung

Die Herstellerin bietet 30 Prozent nach Auftragsbestätigung, 40 Prozent vor Versand und 30 Prozent binnen 15 Tagen nach Lieferung an. Luminara verlangt dagegen 45 Tage ab Lieferung ohne Anzahlung. Keines dieser Modelle ist mit dieser Proforma vereinbart. Vorgesehene Anlieferung ist der 11. Januar 2027; eine bestätigte Bestellung liegt noch nicht vor. Bitte leisten Sie auf dieses Dokument keine Zahlung.

Erstellt von Frauke Eilers, Rechnungswesen. Die Proforma dient der Mengen- und Preisabstimmung und ersetzt weder eine Auftragsbestätigung noch eine Rechnung.''')
    mail(V, '11_2026-09-10_Opening_order_request.eml', 'Marco Rinaldi <m.rinaldi@luminara-strumenti.it>', 'Jan Petersen <j.petersen@weserblick-messtechnik.de>', '2026-09-10T11:39:00', 'LS-PO-0910: quantities for January, payment term still open', '''Dear Jan,

I entered 60 M24 and 15 M48 in the portal this morning, using reference LS-PO-0910 and delivery to our Bologna warehouse on 11 January 2027. The portal shows EUR 35,550 and an advance-payment schedule. It would not let me select 45 days from delivery, so I saved the basket as a request and did not submit a binding order.

Please use the proforma for the quantity discussion only. Giulia has not accepted the payment terms or the direct-account reservation. Our opening quantity is already included in January in the forecast; please do not add it to the twelve-month total a second time.

The six units held under the June loan are still on a separate shelf and are not part of the 75 devices we want to buy. Luca has reported two issues in the loan stock. We will send the stock list and RMA record so your service team can see which units are affected.

Best regards,
Marco Rinaldi
Sales Manager | Luminara Strumenti S.r.l.
Via del Calibratore 22, 40138 Bologna, Italy''')
    csvfile(V, '14_Monthly_forecast_2027.csv', ['Monat', 'M24_Stück', 'M48_Stück', 'Einkauf_M24_EUR', 'Einkauf_M48_EUR', 'Planstatus'], [
        [f'2027-{i:02}', a, b, '420,00', '690,00', 'Erstbevorratung enthalten' if i == 1 else 'unverbindliche Mengenplanung']
        for i, a, b in [(1, 60, 15), (2, 20, 5), (3, 22, 5), (4, 24, 6), (5, 26, 7), (6, 28, 7), (7, 18, 4), (8, 12, 3), (9, 30, 8), (10, 34, 9), (11, 36, 10), (12, 30, 8)]])
    csvfile(V, '15_Bestand_Leihgeraete_2026-09-11.csv', ['Seriennummer', 'Modell', 'Zugang', 'Eigentümerin', 'Ort', 'Zustand', 'Beleg', 'Rückgabe_bis', 'Kaufpreis_EUR'], [
        ['WM24-260601', 'M24', '2026-06-15', 'Weserblick Messtechnik GmbH', 'Bologna Regal 3', 'Vorführung betriebsbereit', 'WM-L-2606', '2026-10-30', '0,00'],
        ['WM24-260602', 'M24', '2026-06-15', 'Weserblick Messtechnik GmbH', 'Bologna Regal 3', 'Vorführung betriebsbereit', 'WM-L-2606', '2026-10-30', '0,00'],
        ['WM24-260603', 'M24', '2026-06-15', 'Weserblick Messtechnik GmbH', 'Bologna RMA-Fach', 'Eingangssignal schwankt', 'RMA-IT-0911-1', '2026-10-30', '0,00'],
        ['WM24-260604', 'M24', '2026-06-15', 'Weserblick Messtechnik GmbH', 'Bologna Regal 3', 'Karton geöffnet, Gerät betriebsbereit', 'WM-L-2606', '2026-10-30', '0,00'],
        ['WM48-260611', 'M48', '2026-06-15', 'Weserblick Messtechnik GmbH', 'Bologna RMA-Fach', 'Klemme beschädigt', 'RMA-IT-0911-2', '2026-10-30', '0,00'],
        ['WM48-260612', 'M48', '2026-06-15', 'Weserblick Messtechnik GmbH', 'Bologna Regal 3', 'Vorführung betriebsbereit', 'WM-L-2606', '2026-10-30', '0,00']])
    doc(V, '16_RMA_Aufnahme_2026-09-11.docx', 'Luminara Strumenti S.r.l. | Luca Serra, Kundendienst', 'RMA Aufnahme für zwei Leihgeräte', 'LS-RMA-0911 | 11. September 2026 | an Ralf Döring, Bremen', '''1 Vorgang RMA-IT-0911-1

Das M24 mit Seriennummer WM24-260603 stammt aus Leihvorgang WM-L-2606, Zugang in Bologna am 15. Juni 2026. Bei einer Vorführung am 9. September schwankte die Anzeige bei konstantem Referenzstrom von 12 mA zwischen 48,9 und 51,2 Prozent. Gemessen wurden am Versorgungsstecker 23,9 V ohne Last und 23,8 V unter Last. Das verwendete Netzgerät ist seit Juni am Vorführtisch im Einsatz. Die anderen M24 zeigten mit demselben Anschluss zwischen 49,9 und 50,1 Prozent.

Luca Serra hat Kabel und Referenzquelle getauscht. Die Schwankung blieb am Gerät. Das Gehäuse wurde nicht geöffnet; die Parametrierdatei wurde lokal gesichert. Der Stecker sitzt fest, äußerlich sind keine Feuchtigkeitsspuren erkennbar. Das Gerät liegt im RMA-Fach und wurde nicht an einen Kunden verkauft. Eine technische Ursache ist noch nicht bestätigt.

2 Vorgang RMA-IT-0911-2

Am M48 mit Seriennummer WM48-260611 ist die Rastnase des Anschlusssteckers beschädigt. Beim Auspacken am 15. Juni wurde keine Beschädigung vermerkt. Marco erinnert sich an einen Sturz des Vorführkoffers beim Termin am 2. September; Luca hat diesen Sturz nicht gesehen. Das Gerät startet, aber der Stecker löst sich bei Zug am Kabel. Es wurde aus der weiteren Vorführung genommen. Ein Austauschstecker ist in Bologna nicht vorhanden.

3 Rücksendung und Kosten

Luca bittet um Rücksendeetiketten und um ein Ersatzgerät für den für 24. September geplanten Kundenbesuch. Der Service in Bremen hat am 11. September telefonisch die beiden Vorgangsnummern vergeben, aber noch keine kostenlose Reparatur oder Ersatzlieferung zugesagt. Luminara hat keine Diagnosepauschale freigegeben. Beide Geräte bleiben Eigentum von Weserblick und dürfen nicht als gekaufte Restbestände verrechnet werden.

4 Gerätezuordnung

Der übrige Leihbestand umfasst WM24-260601, WM24-260602, WM24-260604 und WM48-260612. Die Rückgabe aller sechs Leihgeräte ist nach der Leihbestätigung für den 30. Oktober 2026 vorgesehen. Eine Verlängerung oder Umwidmung zum Servicepool ist bislang nicht bestätigt.

Luca Serra
Kundendienst | l.serra@luminara-strumenti.it''')
    mail(V, '17_2026-09-14_Warranty_and_service.eml', 'Luca Serra <l.serra@luminara-strumenti.it>', 'Ralf Döring <r.doering@weserblick-messtechnik.de>', '2026-09-14T09:05:00', 'RMA cases and the proposed 48-hour exchange service', '''Dear Ralf,

The two devices in our RMA record are loan units, not paid stock. We will not send them until we have the shipping instructions. For the M24 we still see the fluctuating signal after changing the cable and reference source. For the M48 I cannot confirm when the connector was damaged. Please keep the two cases separate when deciding the next step.

For next year's business, Giulia would like to offer a 48-hour replacement service to our customers. I cannot deliver that from four presentation M24 and two presentation M48 that must go back in October. We need a separate pool or we need to charge customers for a different service level. I can check connections and swap devices, but I cannot calibrate the M48 here.

Your contract draft says the dealer margin covers first response. Does that mean only the desk checks, or also driving to a customer and fitting a replacement? Please confirm the intended scope before we quote any service packages. The EUR 95 hourly rate is still our proposal, not an agreed charge.

Best regards,
Luca Serra
Customer Service | Luminara Strumenti S.r.l.
Via del Calibratore 22, 40138 Bologna, Italy''')
    doc(V, '18_Online_Vertriebskonzept_2026-09-14.docx', 'Luminara Strumenti S.r.l. | Marco Rinaldi', 'Onlineverkauf und Markteinführung Italien', 'LS-ONLINE-0914 | 14. September 2026 | an Jan Petersen', '''1 Eigenes Verkaufsangebot

Luminara plant für Januar 2027 einen italienischsprachigen Produktbereich mit M24 und M48. Als Verkäuferin erscheint ausschließlich Luminara mit eigenen Liefer- und Zahlungsbedingungen. Geplant sind Nettoverkaufspreise von 585,00 EUR für M24 und 950,00 EUR für M48. Diese Preise stammen aus unserer Kalkulation vom 7. September und wurden nicht von Weserblick vorgegeben. Rabattangebote an einzelne Kunden werden von Marco freigegeben.

Kunden sollen Angebote anfordern oder Geräte unmittelbar aus dem örtlichen Bestand bestellen können. Die Rechnung wird von Luminara erstellt. Bei einem vom Hersteller direkt belieferten vorbehaltenen Kunden wird kein Geräteverkauf von Luminara ausgewiesen. Ein zusätzlicher Serviceauftrag wird separat angeboten und abgerechnet.

2 Lieferadressen und Portalkonten

Die geplante Shopsoftware kann Lieferadressen außerhalb Italiens annehmen. Wir möchten solche eingehenden Bestellungen nicht automatisch sperren. Aktive Werbung ist zunächst auf Italien ausgerichtet; für Januar bis März sind 9.000 EUR vorgesehen. Offene Marktplätze sind zum Start nicht geplant. Die Forderung nach vorheriger Zustimmung für jede spätere Marktplatznutzung möchten wir vor Vertragsschluss genauer besprechen.

Im Herstellerportal sind Officine Roveta und Vetraria Selce weiterhin mit italienischen Lieferadressen sichtbar. Diese Firmen stehen nicht im Direktkundenverzeichnis als vorbehaltene Konten. Marco möchte, dass ihre neuen Anfragen ab Januar an Luminara gehen. Der von Weserblick genannte Umstellungstermin April würde die ersten drei Monate unseres Lageraufbaus betreffen. Für diese Phase besteht noch keine Abrede zur Umsatzzuordnung.

3 Inhalte und Freigabe

Die italienischen Kurztexte werden von Marco erstellt und von Luca technisch gegengelesen. Messbereich, Versorgung und zulässige Umgebung werden aus dem Herstellerdatenblatt übernommen. Sicherheitsbezogene Übersetzungen gehen vor Veröffentlichung an Ralf Döring. Das Erscheinen als exklusiver Vertriebspartner wird erst nach Unterzeichnung des Vertrags freigegeben. Bis dahin wird kein entsprechender Hinweis veröffentlicht.

Marco Rinaldi
Vertriebsleitung Luminara Strumenti S.r.l.''')
    doc(V, '19_Kundendatenfelder_KD-IT-01.docx', 'Weserblick Messtechnik GmbH | Jan Petersen', 'Monatliche Vertriebsdatei Italien', 'KD-IT-01 V1 | 15. September 2026 | Vorschlag an Luminara', '''1 Übermittlung

Weserblick schlägt eine monatliche UTF-8-Datei vor, die bis zum fünften Arbeitstag des Folgemonats in einen getrennten Uploadbereich gestellt wird. Der Zugriff in Bremen soll auf Exportvertrieb und Serviceleitung beschränkt werden. Der Uploadbereich ist noch nicht eingerichtet. Bis zur Abstimmung sollen keine Kundenlisten per allgemeinem E-Mail-Verteiler versandt werden.

2 Felder für ausgelieferte Geräte

Jede Zeile soll ein Gerät bezeichnen. Vorgesehen sind Luminara-Kundennummer, Kundenfirma, Straße, Postleitzahl, Ort, Name des technischen Ansprechpartners, geschäftliche E-Mail-Adresse, Gerätetyp, Seriennummer, Lieferdatum, Nettoverkaufspreis und Einsatzbereich. Mehrere Geräte eines Auftrags erhalten jeweils eine eigene Zeile. Für reine Lagerlieferungen ohne Endkunden soll die Zeile erst nach dem Wiederverkauf ergänzt werden.

Weserblick möchte die Seriennummern bei technischen Änderungen einem Einsatzort zuordnen und das Mengenbild mit der Planung vergleichen. Der Nettoverkaufspreis soll nach dem Wunsch von Anneke Martens zusätzlich zeigen, in welchen Branchen sich M48 absetzt. Giulia Bellini hat am 3. September die routinemäßige Übermittlung von Verkaufspreisen und persönlichen Kontaktangaben abgelehnt. Der Feldumfang ist deshalb noch nicht freigegeben.

3 Offene Angebote

Für noch nicht ausgelieferte Geschäfte wünscht der Exportvertrieb zusätzlich Kundenfirma, Angebotsdatum, Produktmenge, erwarteten Bestellmonat und Vertriebsansprechpartner. Diese Angaben sollen nicht mit ausgelieferten Seriennummern vermischt werden. Luminara hat bisher nur anonymisierte Mengen je Monat angeboten und möchte offene Kundenangebote nicht vollständig übertragen.

4 Aufbewahrung und Ende der Zusammenarbeit

Weserblick schlägt für Gerätezuordnungen zehn Jahre Aufbewahrung und für offene Angebote zwölf Monate nach dem letzten Kontakt vor. Luminara bittet um eine getrennte Festlegung für technische Rückverfolgung und Vertriebsnutzung sowie eine Benennung der Zugriffsberechtigten. Für das Vertragsende besteht noch keine Einigung, welche Kontakte neben den zu betreuenden Seriennummern übergeben werden. Die Datei wird erst nach gemeinsamer Freigabe verwendet.

Jan Petersen
Exportvertrieb''')
    mail(V, '20_2026-09-16_End_stock.eml', 'Giulia Bellini <g.bellini@luminara-strumenti.it>', 'Anneke Martens <a.martens@weserblick-messtechnik.de>', '2026-09-16T16:12:00', 'Stock at expiry and the six loan units', '''Dear Anneke,

Before committing to annual minimum purchases, I need us to agree what happens to saleable stock when the three-year term ends. We would be buying stock to support your products in Italy, not just ordering against confirmed customer sales. A discretionary buy-back does not give me a number I can put into our cash plan.

Our proposal is repurchase at original net cost for unopened current-model units delivered in the last twelve months, if the contract expires or you end it without a material breach on our side. I am not asking you to buy back damaged customer returns as new goods. Older saleable units need a twelve-month sell-off period, including our web shop.

The six June loan units are different. We have not paid for them and they remain yours. We will return them under the loan terms unless the service team agrees another arrangement. Please do not treat those six units as evidence that we already have a paid exchange pool. The opening order of 75 new devices is still unconfirmed.

Best regards,
Giulia Bellini
Managing Director | Luminara Strumenti S.r.l.
Via del Calibratore 22, 40138 Bologna, Italy''')
    mail(V, '21_2026-09-17_Finance_terms.eml', 'Frauke Eilers <f.eilers@weserblick-messtechnik.de>', 'Giulia Bellini <g.bellini@luminara-strumenti.it>', '2026-09-17T10:28:00', 'Proforma PF-IT-260909 and requested credit terms', '''Dear Ms Bellini,

The proforma remains EUR 35,550 for 60 M24 and 15 M48. The zero-percent line is conditional on the required evidence and our final invoice review. We have not entered a tax identification number on this document. Please send the company and delivery records to the agreed accounts contact; the sales portal entry alone is not sufficient for release.

For payment, our offer still contains EUR 10,665 after order confirmation, EUR 14,220 before dispatch and EUR 10,665 within 15 days after delivery. I have recorded your request for one payment 45 days after delivery and a EUR 60,000 credit limit, but neither has been approved. Assuming delivery on 11 January 2027, your proposed due date of 25 February is arithmetically correct.

No payment is due against the proforma, and no invoice for this opening order has been issued. Please wait for a jointly agreed payment schedule and the final order confirmation before transferring funds.

Kind regards,
Frauke Eilers
Accounts | Weserblick Messtechnik GmbH
Am Speicherbogen 18, 28197 Bremen, Germany''')
    mail(V, '22_2026-09-18_Daten_und_Portal.eml', 'Jan Petersen <j.petersen@weserblick-messtechnik.de>', 'Anneke Martens <a.martens@weserblick-messtechnik.de>', '2026-09-18T13:46:00', 'Datenfelder und italienische Portalbestellungen', '''Hallo Anneke,

Giulia gibt uns monatlich Mengen, Seriennummern, Kundenfirmen, Lieferorte und Einsatzbereiche. Bei persönlichen Ansprechpartnern will sie auf konkrete Servicefälle abstellen. Ihre offenen Angebote und Wiederverkaufspreise möchte sie nicht fortlaufend hochladen. Sie befürchtet, dass wir daraus eigene Direktangebote machen. Ich habe ihr dazu keine Zusage gegeben und die Uploadfreigabe zurückgehalten.

Die Portalumstellung lässt sich nach heutigem Stand erst im April einrichten. Für Januar bis März könnten wir Anfragen manuell weiterreichen. Officine Roveta mit acht M24 und Vetraria Selce mit vier M48 sind noch nicht bestätigt. Beide stehen außerhalb der Adria-Ausnahme. Wenn wir die Aufträge selbst bestätigen, brauchen wir mit Luminara eine klare Verständigung über diese Übergangsmonate.

Der Forecast umfasst 340 M24 und 87 M48. Zu unseren Einkaufspreisen sind das 202.830 EUR netto, einschließlich der 75 Geräte im Januar. Das ist eine Planung oberhalb beider bisheriger Mindestbezugsangebote, keine feste Bestellung. Ich würde daraus noch keine Bereitschaft zu unserem 180.000-EUR-Vertragsbetrag ableiten.

Viele Grüße
Jan Petersen
Exportvertrieb | Weserblick Messtechnik GmbH
Am Speicherbogen 18, 28197 Bremen''')
    doc(V, '23_Logistik_Leihbestand_2026-09-18.docx', 'Weserblick Messtechnik GmbH | Ralf Döring / Frauke Eilers', 'Abstimmung zur Rückführung des Leihbestands', 'WM-L-2606 Rückführung | 18. September 2026 | an Luca Serra', '''1 Bestand und Eigentum

Unsere Leihbestätigung WM-L-2606 vom 10. Juni 2026 betrifft vier M24 mit den Seriennummern WM24-260601 bis WM24-260604 und zwei M48 mit den Seriennummern WM48-260611 und WM48-260612. Der Warenausgang war am 12. Juni, der Zugang in Bologna am 15. Juni. Die Geräte wurden zur Vorführung überlassen; ein Kaufpreis wurde nicht berechnet. Rückgabe ist bis 30. Oktober 2026 vorgesehen.

2 Zwei vorgezogene Rücksendungen

Für WM24-260603 und WM48-260611 werden getrennte RMA-Etiketten erstellt. Die Geräte sollen mit dem jeweils benutzten Anschlussstecker und dem zugehörigen Prüfblatt zurückkommen. Bitte nicht gemeinsam in einem unbeschrifteten Innenkarton verpacken. Das M24 wird wegen der gemeldeten Signalabweichung geprüft; beim M48 wird zunächst der Steckerschaden aufgenommen. Eine Übernahme der Reparaturkosten ist noch nicht zugesagt. Die Erstfracht für diese beiden Leihgeräte organisiert Weserblick.

3 Weiterer Vorführtermin

Für den Besuch am 24. September kann Luminara eines der drei übrigen betriebsbereiten M24 und das M48 WM48-260612 verwenden. Ein zusätzliches Ersatzgerät ist derzeit nicht disponiert. Die verbleibenden vier Geräte sind anschließend bis zum Rückgabetermin gesammelt zurückzugeben, sofern keine schriftliche Verlängerung erfolgt.

4 Trennung von der Erstbestellung

Der geplante Neukauf von 60 M24 und 15 M48 ist nicht zur Verrechnung mit dem Leihbestand vorgesehen. Die Preis- und Zahlungsabstimmung hierzu läuft über Jan Petersen und Frauke Eilers. Für die sechs Leihgeräte besteht weder eine offene Kaufpreisforderung noch eine erteilte Gutschrift. Eine Umwidmung in einen Austauschpool für 2027 wurde noch nicht beschlossen.

Ralf Döring, Serviceleitung
Frauke Eilers, Rechnungswesen''')
    doc(V, '24_Gespraechsnotiz_2026-09-22.docx', 'Weserblick Messtechnik GmbH | Jan Petersen', 'Telefonat mit Luminara am 22 September', 'VP-0922 | 22. September 2026, 10:00 bis 10:52 Uhr | an Anneke Martens und Giulia Bellini', '''1 Teilnehmer und Bestellstand

Teilgenommen haben Anneke Martens und Jan Petersen für Weserblick sowie Giulia Bellini und Marco Rinaldi für Luminara. Die vorgesehenen Einkaufspreise von 420,00 EUR für M24 und 690,00 EUR für M48 wurden für Lieferungen 2027 erneut genannt. Die Mengenanfrage über 60 M24 und 15 M48 bleibt gespeichert, ist aber nicht bestätigt. Giulia hält an 45 Tagen ab Lieferung fest; Anneke will dazu erst die Freigabe der Buchhaltung abwarten. Eine Zahlung ist nicht erfolgt.

2 Gebiet und Direktkunden

Italien bleibt das besprochene Gebiet, die Laufzeit soll vom 1. Januar 2027 bis 31. Dezember 2029 reichen. Anneke möchte Adria Cartoni einschließlich zukünftiger Standorte unmittelbar bedienen. Giulia würde nur Parma und Verona unter dem vorhandenen Rahmenvertrag ausnehmen. Über Bari und zukünftige Konzernunternehmen gab es keine Verständigung. Auch die Behandlung europaweiter Rahmenverträge wurde nicht abschließend besprochen.

Jan erläuterte den technischen Portaltermin April 2027. Giulia möchte vor Januar eine manuelle Weiterleitung für nicht vorbehaltene italienische Kunden. Anneke hat keine Anrechnung eigener Direktumsätze auf Luminaras Mindestbezug zugesagt. Die beiden bestehenden Portalanfragen bleiben vorläufig unbestätigt. Luminara möchte unaufgeforderte Bestellungen aus dem Ausland im eigenen Shop weiter annehmen können.

3 Mindestbezug und Kalkulation

Die zwölf Forecast-Zeilen ergeben 340 M24 und 87 M48, somit 202.830,00 EUR geplanten Nettobezug. Die Erstbevorratung ist in Januar enthalten. Giulia betonte, dass die geplanten Mengen keine jährliche Abnahmeverpflichtung von 180.000,00 EUR bestätigen. Luminara bietet weiterhin 150.000,00 EUR im ersten Jahr, 190.000,00 EUR im zweiten und 230.000,00 EUR im dritten Jahr. Weserblick bleibt vorerst bei 180.000,00 EUR, 220.000,00 EUR und 260.000,00 EUR.

4 Service, Daten und Restbestand

Die Beteiligten trennten erneut die Handelsspanne aus Eigenverkäufen von gesondert beauftragten Besuchen bei Direktkunden. Ein Stundenpreis und ein Austauschpool sind noch nicht vereinbart. Der Geräteverbleib aus der Juni-Leihe folgt weiterhin der Rückführung vom 18. September. Die beiden RMA-Fälle sind offen.

Giulia hält an einem verbindlichen Rückkauf jüngerer ungeöffneter Ware bei regulärem Ablauf fest. Anneke möchte nur einzelne Rücknahmen freigeben. Die monatliche Weitergabe persönlicher Ansprechpartner, Wiederverkaufspreise und offener Angebote bleibt ebenso offen wie eine Verwendung der Daten nach Vertragsende.

5 Vertragsfassung

Der deutsche Entwurf vom 20. August und Luminaras englischer Gegenvorschlag vom 3. September sind unverändert nebeneinander im Umlauf. Deutsches Recht mit Bremen und italienisches Recht mit Bologna wurden als jeweilige Vorschläge genannt; eine Auswahl wurde nicht getroffen. Auch eine maßgebliche Sprachfassung wurde nicht vereinbart. Keine Seite hat einen Vertriebsvertrag unterschrieben oder eine Exklusivität als bereits wirksam bestätigt.

Jan Petersen
Versandt an die Teilnehmer am 22. September 2026 um 14:10 Uhr.''')


def font(size, bold=False):
    return screen_font(size, bold)


def drawing_text(draw, xy, text, size=23, bold=False, fill='#202528'):
    draw.multiline_text(xy, text, font=font(size, bold), fill=fill, spacing=7)


def layout():
    image = Image.new('RGB', (1800, 1400), 'white')
    draw = ImageDraw.Draw(image)
    drawing_text(draw, (65, 42), 'Wertau Lebensmittel GmbH   |   Linie 3, Augsburg', 34, True)
    drawing_text(draw, (65, 95), 'WA-L3-04   ·   08.09.2026   ·   Paula König   ·   Aufstellung zur Abstimmung', 23)
    ox, oy, scale = 150, 1070, 100

    def pt(x, y):
        return (int(ox + x * scale), int(oy - y * scale))

    def box(x1, y1, x2, y2, color, label, small=22):
        left, top = pt(x1, y2)
        right, bottom = pt(x2, y1)
        draw.rectangle((left, top, right, bottom), fill=color, outline='#42494e', width=3)
        drawing_text(draw, (left + 12, top + 15), label, small)

    draw.rectangle((*pt(0, 8), *pt(10, 0)), outline='#24292d', width=7)
    for x in range(11):
        a, b = pt(x, 0), pt(x, 8)
        draw.line((a, b), fill='#ebedef', width=1)
        drawing_text(draw, (a[0] - 8, oy + 18), str(x), 18)
    for y in range(9):
        a, b = pt(0, y), pt(10, y)
        draw.line((a, b), fill='#ebedef', width=1)
        drawing_text(draw, (ox - 42, a[1] - 10), str(y), 18)
    drawing_text(draw, (140, 208), 'Nordwand', 23, True)
    drawing_text(draw, (1040, 1120), 'x in Metern', 19)
    drawing_text(draw, (55, 225), 'y [m]', 20)
    box(0.5, 4.4, 1.8, 5.8, '#e4e8ed', 'Linien-\nverteiler\nE1', 21)
    box(4, 4.5, 10, 7.5, '#fff3d1', 'Bereitstellfläche Nord\n6,0 x 3,0 m = 18 m²\nRäumung zum 12.07.2027\nnoch nicht bestätigt.', 23)
    draw.line((*pt(7, 8), *pt(9.4, 8)), fill='white', width=9)
    drawing_text(draw, pt(7, 8.55), 'Nordtor 2,40 m', 21)
    draw.line((*pt(0, .5), *pt(0, 2.6)), fill='white', width=9)
    drawing_text(draw, (24, 915), 'Westtor\n2,10 m', 19)
    drawing_text(draw, (1380, 270), 'N', 26, True)
    draw.line((1394, 360, 1394, 310), fill='#202528', width=4)
    draw.polygon([(1394, 298), (1385, 315), (1403, 315)], fill='#202528')
    box(.8, 1.6, 3.2, 3.6, '#d8ece9', 'Zweifachdosierer\nHP-D2\n2,40 x 2,00 m', 21)
    box(3.2, 2.25, 4.2, 2.95, '#d8ece9', '', 17)
    drawing_text(draw, pt(3.2, 3.5), 'Puffer\n28 Becher', 17)
    box(4.2, 1.6, 5.7, 3.6, '#e6e7e9', 'Versiegler\nBestand', 21)
    box(5.7, 1.85, 6.3, 3.35, '#e6e7e9', 'Etik.', 19)
    box(6.3, 1.6, 9.2, 3.6, '#d8ece9', 'Kamera / Ausschleuser\nGutlauf bis T4', 21)
    box(7.35, 1.8, 8.15, 2.2, '#f3c9c9', 'Box', 18)
    draw.line((*pt(.8, 2.6), *pt(8.8, 2.6)), fill='#27766c', width=7)
    draw.polygon([pt(8.8, 2.6), pt(8.6, 2.7), pt(8.6, 2.5)], fill='#27766c')
    for n, x in [('T1', .8), ('T2', 4.2), ('T3', 6.3), ('T4', 8.8)]:
        xx, yy = pt(x, 2.6)
        draw.ellipse((xx-6, yy-6, xx+6, yy+6), fill='#202528')
        drawing_text(draw, (xx-13, yy+28), n, 18, True)
    drawing_text(draw, pt(1, 1.35), 'Bediengang Süd   1,20 m frei halten', 23)
    drawing_text(draw, pt(2, 4.08), 'Wartungsraum Nord   0,90 m', 20)
    drawing_text(draw, pt(2.5, 5.7), 'D1\n6 bar', 19)
    drawing_text(draw, pt(.12, 3.3), 'P1', 18, True)
    drawing_text(draw, pt(2.2, .5), 'R1', 18, True)
    drawing_text(draw, (1400, 450), 'Förderhöhe\n920 mm\n\nGrün: Helix\nGrau: Bestand\nGelb: Logistik', 23)
    drawing_text(draw, (1400, 780), 'E1: 400 V / 125 A\nD1: 800 Nl/min\nP1: Produkt DN 40\nR1: Bodenablauf', 20)
    drawing_text(draw, (150, 1190), 'Förderrichtung links nach rechts. Koordinatenursprung: südwestliche Hallenecke.', 22)
    drawing_text(draw, (150, 1230), 'Breiteste Transportbaugruppe: 2,25 m. Nordzufahrt: 2,40 m. Westtor: 2,10 m.', 22)
    drawing_text(draw, (150, 1270), 'Keine Montagefreigabe. Maße vor Fertigungsbeginn am Bestand prüfen.', 22)
    image.save(remember(P, '15_Stellplan_WA-L3-04.png'))


def portal():
    image = Image.new('RGB', (1600, 1130), '#f5f6f8')
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 84), fill='#28383e')
    drawing_text(draw, (45, 24), 'WESERBLICK   Partnerportal', 30, True, 'white')
    drawing_text(draw, (1050, 29), 'Luminara Strumenti S.r.l.', 24, False, 'white')
    drawing_text(draw, (65, 126), 'Anfragen / LS-PO-0910', 26)
    drawing_text(draw, (65, 181), 'Erstbevorratung Bologna', 36, True)
    drawing_text(draw, (65, 240), 'Gespeichert am 10.09.2026 um 10:54   |   Marco Rinaldi', 23)
    draw.rectangle((65, 290, 1535, 355), fill='#fff0c2')
    drawing_text(draw, (85, 308), 'Anfrage gespeichert. Noch keine Bestellung übermittelt.', 26, True)
    drawing_text(draw, (65, 393), 'Lieferstelle: Via del Calibratore 22, 40138 Bologna, Italien', 26)
    drawing_text(draw, (65, 438), 'Wunschtermin: 11.01.2027    |    Konditionen: WM-IT-260714', 24)
    draw.rectangle((65, 507, 1535, 566), fill='#dce4e8')
    for x, label in [(88, 'Produkt'), (730, 'Menge'), (970, 'Einzelpreis netto'), (1280, 'Wert netto')]:
        drawing_text(draw, (x, 524), label, 23, True)
    for y, values in [(594, ['Messmodul M24', '60', '420,00 EUR', '25.200,00 EUR']), (668, ['Messmodul M48', '15', '690,00 EUR', '10.350,00 EUR'])]:
        draw.rectangle((65, y-10, 1535, y+51), fill='white')
        for x, value in zip([88, 740, 995, 1280], values):
            drawing_text(draw, (x, y), value, 25)
    drawing_text(draw, (990, 753), 'Warenwert    35.550,00 EUR', 28, True)
    drawing_text(draw, (65, 821), 'Angebotsmodell: 30 % nach Bestätigung / 40 % vor Versand / 30 % nach Lieferung', 24)
    drawing_text(draw, (65, 865), 'Wunsch des Kunden: 45 Tage ab Lieferung. Freigabe ausstehend.', 24)
    drawing_text(draw, (65, 909), 'Steueransatz 0 % nur unter Nachweisvorbehalt gemäß Proforma PF-IT-260909.', 23)
    draw.rectangle((65, 981, 465, 1045), outline='#66777d', width=2)
    drawing_text(draw, (103, 1000), 'Anfrage bearbeiten', 25)
    draw.rectangle((500, 981, 1010, 1045), fill='#e1e4e7')
    drawing_text(draw, (525, 1000), 'Bestellung noch nicht freigegeben', 25, False, '#60676b')
    drawing_text(draw, (65, 1080), 'Die Anfrage reserviert keinen Liefertermin. Eine Auftragsbestätigung liegt nicht vor.', 21)
    image.save(remember(V, '12_Portal_Anfrage_2026-09-10.png'))


def readme(folder, title, summary):
    names = sorted(FILES[folder])
    assert len(names) == 24, (folder.name, len(names))
    assert len(set(names)) == 24
    for n in names:
        assert (folder / n).is_file(), n
    release = 'https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download'
    text = f'''# {title}

## 1 Vorgang

{summary} Bearbeitungsstand: 22. September 2026. Zugeordnetes Plugin: `corporate-contract-law`. Veröffentlichung mit Version 445.1.0.

## 2 Downloads

Die drei Downloadfassungen werden im zentralen Release erstellt. Das Originalformat-Archiv enthält die nativen Quellen; das Einzel-PDF-Archiv enthält jede Unterlage separat. Die Links auf Release-Archive sind erst nach Veröffentlichung dieser Akte verfügbar.

> {NOTICE[0]}
>
> {NOTICE[1]}

| Fassung | Download |
| --- | --- |
| Gesamt-PDF | [Akte als Gesamt-PDF](gesamt-pdf/{folder.name}_gesamt.pdf) |
| Akten-ZIP | [Native Originaldateien]({release}/testakte-{folder.name}.zip) |
| Einzel-PDF-ZIP | [Unterlagen als Einzel-PDFs]({release}/testakte-{folder.name}-einzelpdfs.zip) |

## 3 Quellen

24 getrennte Quellen. Vertragsfassungen, Nachrichten und Rechenstände tragen jeweils ihre eigene Herkunft und ihren eigenen Bearbeitungsstand. Die Umsetzung im Jahr 2027 ist geplant; ein Abschluss des Hauptvertrags wird nicht vorweggenommen.

'''
    text += '\n'.join(f'- [{n}]({n})' for n in names)
    text += '\n\n## 4 Zuordnung\n\n[Corporate Contract Law](../../corporate-contract-law/README.md) · [Aktenübersicht](../README.md)\n'
    (folder / 'README.md').write_text(text, encoding='utf-8')


def verify():
    sys.dont_write_bytecode = True
    import importlib.util
    spec = importlib.util.spec_from_file_location('document_quality', ROOT / 'scripts/validate-testakten-dokumentqualitaet.py')
    quality = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(quality)
    for folder in (P, V):
        files = sorted(p for p in folder.iterdir() if p.is_file() and re.match(r'^\d{2}_', p.name))
        assert len(files) == 24, (folder, len(files))
        assert [p.name[:2] for p in files] == [f'{n:02}' for n in range(1, 25)]
        for path in files:
            if path.suffix == '.png':
                Image.open(path).verify()
                continue
            body = quality.export_text(path)
            assert not quality.SYNTHETIC_EMAIL_PATTERN.search(body), path
            assert not any(m in body.lower() for m in ['testakte', 'fiktiv', 'platzhalter', 'formathinweis', 'kartellverstoß', '89b analog']), path
            assert '\u00a7' not in body, path
            for label, pattern in quality.EXPORT_META_PATTERNS.items():
                assert not pattern.search(body), (path, label)
            if path.suffix in {'.docx', '.eml', '.txt'}:
                assert len(body.strip()) >= quality.MIN_FORMAL_TEXT, (path, len(body.strip()))
            if path.suffix == '.docx':
                assert quality.is_a4(Document(path)), path
                assert not quality.language_prose_errors(body, path, Document(path).core_properties.language), (path, quality.language_prose_errors(body, path, Document(path).core_properties.language))
            if path.suffix == '.eml':
                assert not quality.eml_quality_errors(path), quality.eml_quality_errors(path)
                language = str(quality.eml_message(path).get('Content-Language', ''))
                assert not quality.language_prose_errors(body, path, language), (path, quality.language_prose_errors(body, path, language))
        print(f'{folder.name}: 24 native Quellen; Text, E-Mail-Köpfe, Domains und A4 geprüft.')


def render_documents(qa):
    renderer = shutil.which('pdftoppm')
    if not renderer:
        raise RuntimeError('Für die Sichtprüfung fehlt pdftoppm. Poppler installieren.')
    sources = [source for folder in (P, V) for source in sorted(folder.glob('*.docx'))]
    pdfs = render_office_batch(sources)
    missing = [str(source) for source in sources if source not in pdfs]
    if missing:
        raise RuntimeError('PDF-Konvertierung fehlt; LibreOffice oder SOFFICE prüfen: ' + ', '.join(missing))
    qa.mkdir(parents=True, exist_ok=True)
    for source in sources:
        out = qa / source.parent.name / source.stem
        out.mkdir(parents=True, exist_ok=True)
        pdf = out / (source.stem + '.pdf')
        pdf.write_bytes(pdfs[source])
        subprocess.run([renderer, '-scale-to-x', '1100', '-scale-to-y', '1556', '-png', str(pdf), str(out / 'page')], check=True)
    pages = sorted(qa.glob('corporate-*/*/page-*.png'))
    for offset in range(0, len(pages), 12):
        board = Image.new('RGB', (1800, 2120), '#e7e7e7')
        canvas = ImageDraw.Draw(board)
        for j, page in enumerate(pages[offset:offset+12]):
            tile = Image.open(page).convert('RGB')
            tile.thumbnail((430, 610))
            x, y = (j % 4)*450, (j//4)*700
            board.paste(tile, (x+10, y+45))
            label = page.parent.name[:31] + '\n' + page.name
            drawing_text(canvas, (x+10, y+3), label, 15)
        board.save(qa / f'contact-{offset//12+1:02}.png')
    print(f'{len(pages)} Dokumentseiten gerendert: {qa}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true', help='Prüft nur vorhandene Quellen.')
    parser.add_argument('--render', type=Path, help='Rendert DOCX-Seiten zur visuellen Kontrolle, ohne PDF-Endexport.')
    args = parser.parse_args()
    if args.check:
        verify()
    elif args.render:
        render_documents(args.render)
    else:
        runtime = node_binary()
        workbook = str(ROOT / 'scripts/build-corporate-projekt-vertrieb-workbooks.mjs')
        environment = dict(os.environ, AKTEN_PYTHON=sys.executable)
        subprocess.run([runtime, workbook, '--check-runtime'], env=environment, check=True)
        screen_font(12)
        screen_font(12, bold=True)
        project()
        distribution()
        layout()
        portal()
        for folder, name in [(P, '17_Meilensteinzahlungen_und_Termine.xlsx'), (V, '13_Haendlerkalkulation_2027.xlsx')]:
            FILES[folder].append(name)
        subprocess.run([runtime, workbook], env=environment, check=True)
        readme(P, 'Dosier- und Verpackungsintegration in Augsburg', 'Wertau Lebensmittel GmbH und Helix Prozessautomation GmbH verhandeln über die Integration der Linie 3 zum angebotenen Preis von 675.000 EUR netto. Vorplanung, Leistungsumfang, Anschlussdaten und Vertragsfassungen liegen getrennt vor.')
        readme(V, 'Messtechnikvertrieb in Italien', 'Weserblick Messtechnik GmbH aus Bremen und Luminara Strumenti S.r.l. aus Bologna verhandeln über einen dreijährigen Vertrieb der Geräte M24 und M48 in Italien. Deutscher Herstellerentwurf und englischer Gegenvorschlag stehen neben Angeboten, Korrespondenz und Vertriebszahlen.')
        verify()


if __name__ == '__main__':
    main()
