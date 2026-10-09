#!/usr/bin/env python3
"""Bauvergabe: Installationsgrenzen, echte Anlagen und durchgehende Fallzahlen."""
from datetime import date, datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import unquote
import zipfile
from docx import Document
from openpyxl import load_workbook
from bauvergabe_falldaten import CASES, PLUGINS, STAGES, total
from bauvergabe_aktentexte import documents
from prompt_profiles import validate_files
from quality_lab import validate_profile
from release_routing import validate_plugin_version
from testakte_zip_common import working_dump_archive_pairs, preserves_directories
from testakte_disclaimer import NOTICE_BYTES
from testakte_einzelpdf_common import document_arcname_pairs

ROOT=Path(__file__).resolve().parents[1]
NATIVE_COUNTS={'bauvergabe-klinikum-muenster':38,'bauvergabe-wohnhaus-bielefeld':39}
NATIVE_EXTENSIONS={'.docx','.pdf','.eml','.csv','.xlsx'}

def record_date(item):
    value=item['date']
    return datetime.strptime(value,'%d.%m.%Y').date() if '.' in value else datetime.fromisoformat(value).date()

def expected_native_names(case):
    """Unabhängig vom Exportfilter: erzeugte Aktenstücke und explizite Rechenbelege."""
    records=documents(case)
    names={item['file'] for item in records}
    names.update(str(Path(item['file']).with_suffix('.pdf')) for item in records if item.get('pdf'))
    names.update({
        '01-vergabeunterlagen/07_Mengen_Preisblatt.csv',
        '02-vergabeverfahren/04_Angebotseingang.csv',
        '03-bieterarbeit/02_Angebotskalkulation.xlsx',
        '05-nachtragsmanagement/04_Aufmass_N01.csv',
        '05-nachtragsmanagement/06_Nachtragskosten.xlsx',
    })
    return names

def module(name, filename):
    spec=importlib.util.spec_from_file_location(name,ROOT/'scripts'/filename)
    mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod);return mod

class Bauvergabe(unittest.TestCase):
    def test_five_nested_installations_and_no_sixth_wrapper(self):
        market=json.loads((ROOT/'.claude-plugin/marketplace.json').read_text())
        entries={x['name']:x for x in market['plugins']}
        self.assertNotIn('bauvergabe',entries)
        self.assertFalse((ROOT/'bauvergabe/.claude-plugin/plugin.json').exists())
        for slug in PLUGINS:
            with self.subTest(plugin=slug):
                self.assertEqual(entries[slug]['source'],'./bauvergabe/'+slug)
                validate_plugin_version(entries[slug], market['version'], root=ROOT)
                p=ROOT/entries[slug]['source']
                skills=list((p/'skills').glob('*/SKILL.md'))
                self.assertEqual(len(skills),11)
                for kind in ('.claude-plugin','.codex-plugin'):
                    m=json.loads((p/kind/'plugin.json').read_text())
                    self.assertEqual(m['name'],slug)
                    self.assertEqual(m['version'],entries[slug]['version'])
                self.assertEqual(validate_files(p,slug,ROOT),[])
                profile=json.loads((ROOT/'quality/evals'/f'{slug}.json').read_text())
                validate_profile(profile,slug,p,ROOT)
                self.assertEqual({x['target_skill'] for x in profile['cases']},{x.parent.name for x in skills})
                for kind,key in [('werkstatt','workshop_review'),('schnellstart','mini_review')]:
                    b=(p/f'{slug}-{kind}.md').read_bytes()
                    self.assertEqual(hashlib.sha256(b).hexdigest(),profile[key]['sha256'])
                    if kind=='schnellstart':self.assertLessEqual(len(b),7500)
                for skill in skills:
                    for target in re.findall(r'\]\(([^\s)]+)\)',skill.read_text()):
                        if '://' not in target and not target.startswith('#'):
                            dest=(skill.parent/unquote(target.split('#')[0])).resolve()
                            self.assertTrue(dest.is_relative_to(p.resolve()),f'Nicht installierter Skillpfad: {skill}: {target}')
                            self.assertTrue(dest.exists(),f'Fehlende Ressource: {skill}: {target}')
                            self.assertNotRegex(dest.name,r'-(?:werkstatt|schnellstart)\.')

    def test_discovery_uses_source_not_only_root_directory(self):
        generator=module('bau_prompt_generator','generate-werkstatt-und-schnellstart-prompts.py')
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);(root/'.claude-plugin').mkdir()
            nested=root/'familie/fachplugin/.claude-plugin';nested.mkdir(parents=True)
            (nested/'plugin.json').write_text('{}')
            (root/'.claude-plugin/marketplace.json').write_text(json.dumps({'plugins':[{'name':'fachplugin','source':'./familie/fachplugin'}]}))
            with patch.object(generator,'REPO',root):
                self.assertEqual(generator.plugin_dirs(),[nested.parent.resolve()])

    def test_discovery_deduplicates_and_ignores_unusable_marketplace_sources(self):
        generator=module('bau_prompt_generator_edge','generate-werkstatt-und-schnellstart-prompts.py')
        with tempfile.TemporaryDirectory() as t:
            root=Path(t).resolve();(root/'.claude-plugin').mkdir()
            wanted=[]
            for relative in ('familie/fachplugin','flachplugin','gerichtsplugins/gericht'):
                folder=root/relative;(folder/'.claude-plugin').mkdir(parents=True)
                (folder/'.claude-plugin/plugin.json').write_text('{}');wanted.append(folder)
            outside=root/'outside/.claude-plugin';outside.mkdir(parents=True);(outside/'plugin.json').write_text('{}')
            entries=[{'name':'fachplugin','source':'./familie/fachplugin'},
                     {'name':'alias','source':'./familie/fachplugin'},
                     {'name':'flachplugin','source':'./flachplugin'},
                     {'name':'gericht','source':'./gerichtsplugins/gericht'},
                     {'name':'missing','source':'./nicht-vorhanden'},
                     {'name':'remote','source':{'source':'github','repo':'some/repo'}},
                     {'name':'url','source':'https://example.invalid/repo'},
                     {'name':'escape','source':'./../escaped-plugin'}]
            # A top-level legacy plugin is intentionally still discoverable.
            wanted.append(outside.parent)
            (root/'.claude-plugin/marketplace.json').write_text(json.dumps({'plugins':entries}))
            with patch.object(generator,'REPO',root):
                found=generator.plugin_dirs()
            self.assertEqual(set(found),set(wanted));self.assertEqual(len(found),len(wanted))
        real=generator.plugin_dirs()
        for slug in PLUGINS:self.assertEqual(real.count((ROOT/'bauvergabe'/slug).resolve()),1)

    def test_bid_and_change_amounts_across_native_sources(self):
        for c,want,change in zip(CASES,['7823200.00','1775620.00'],['244532.00','64622.44']):
            with self.subTest(case=c['slug']):
                self.assertEqual(total(c),Decimal(want))
                folder=ROOT/'testakten'/c['slug']
                wb=load_workbook(folder/'03-bieterarbeit/02_Angebotskalkulation.xlsx',data_only=True)
                self.assertEqual(Decimal(str(wb.active['F32'].value)),Decimal(want));wb.close()
                wb=load_workbook(folder/'05-nachtragsmanagement/06_Nachtragskosten.xlsx',data_only=True)
                self.assertEqual(Decimal(str(wb.active['F25'].value)),Decimal(change));wb.close()
                offer=Document(folder/'03-bieterarbeit/04_Angebotsschreiben.docx')
                body='\n'.join(p.text for p in offer.paragraphs)
                self.assertIn(c['winner'],body)
                from bauvergabe_falldaten import money
                self.assertIn(money(total(c)),body)

    def test_case_sequence_and_evidence_are_not_future_backdated(self):
        for c in CASES:
            d=lambda k:date.fromisoformat(c[k])
            self.assertGreaterEqual((d('original_deadline')-d('published')).days,35)
            self.assertGreater(d('deadline'),d('original_deadline'))
            self.assertLess(d('notice'),d('objection'))
            self.assertLessEqual((d('objection')-d('notice')).days,10)
            self.assertLess(d('refusal'),d('petition'))
            self.assertLessEqual((d('petition')-d('refusal')).days,15)
            self.assertLess(d('service'),d('first_award'))
            self.assertLess(d('withdrawal'),d('new_notice'))
            self.assertGreaterEqual((d('first_award')-d('notice')).days,11)
            self.assertGreaterEqual((d('award')-d('new_notice')).days,11)
            self.assertLessEqual(d('award'),date(2027,1,31))
            self.assertLess(d('award'),d('start'))
            source=documents(c)
            self.assertEqual(len(source),len({i['file'] for i in source}))
            for i in source:
                if i['file'].endswith('.docx'):
                    self.assertTrue((ROOT/'testakten'/c['slug']/i['file']).exists())
            by_name={Path(i['file']).stem:i for i in source}
            self.assertEqual((d('deadline')-record_date(by_name['04_Angebotsschreiben'])).days,1)
            self.assertLess(d('deadline'),record_date(by_name['05_Nachforderung']))
            self.assertLess(record_date(by_name['05_Nachforderung']),record_date(by_name['06_Antwort_Aufklaerung']))
            self.assertLess(record_date(by_name['06_Antwort_Aufklaerung']),d('notice'))
            self.assertGreaterEqual(record_date(by_name['03_Nachtragsangebot']),record_date(by_name['05_Bauablauf_Protokoll']))
            if c['short']=='Wohnhaus':
                commitment=next(i for i in source if 'Verpflichtung_Knuff' in i['file'])
                self.assertEqual(record_date(commitment),date(2026,11,20))
                self.assertNotRegex(commitment['body'],r'4\. Dezember|04\.12\.|ersten Upload|Nachreichung')
                self.assertEqual(by_name['06_Antwort_Aufklaerung']['attachments'],[commitment['file']])

    def assert_native_inventory(self,folder,wanted):
        disk={p.relative_to(folder).as_posix() for stage in STAGES for p in (folder/stage).rglob('*')
              if p.is_file() and p.suffix.lower() in NATIVE_EXTENSIONS}
        self.assertEqual(disk,wanted,'Dateibaum enthält fehlende oder unerwartete native Aktenstücke.')
        exported=working_dump_archive_pairs(folder,include_gesamt_pdf=False)
        self.assertEqual({p.relative_to(folder).as_posix() for p,_ in exported},wanted,
                         'Exportfilter hat native Aktenstücke verworfen oder zusätzliche Dateien aufgenommen.')
        self.assertEqual({name for _,name in exported},wanted)
        self.assertEqual(len(exported),len(wanted))

    def test_all_77_native_originals_survive_export_filter(self):
        count=0
        for c in CASES:
            with self.subTest(case=c['slug']):
                wanted=expected_native_names(c);self.assertEqual(len(wanted),NATIVE_COUNTS[c['slug']])
                self.assert_native_inventory(ROOT/'testakten'/c['slug'],wanted);count+=len(wanted)
        self.assertEqual(count,77)

    def test_native_inventory_check_rejects_a_silently_filtered_xlsx(self):
        c=CASES[0];folder=ROOT/'testakten'/c['slug'];wanted=expected_native_names(c)
        filtered=[(folder/name,name) for name in wanted if not name.endswith('02_Angebotskalkulation.xlsx')]
        with patch(__name__+'.working_dump_archive_pairs',return_value=filtered):
            with self.assertRaises(AssertionError):self.assert_native_inventory(folder,wanted)

    def test_combined_pdf_builder_rejects_a_silently_filtered_xlsx(self):
        builder=module('bau_pdf_source_guard','build-bauvergabe-pakete.py')
        for c in CASES:
            folder=ROOT/'testakten'/c['slug']
            self.assertEqual({p.relative_to(folder).as_posix() for p in builder.sources(folder)},expected_native_names(c))
            selected=document_arcname_pairs(folder)
            without_xlsx=[(p,name) for p,name in selected if p.suffix!='.xlsx']
            with patch.object(builder,'document_arcname_pairs',return_value=without_xlsx):
                with self.assertRaisesRegex(RuntimeError,'PDF-Quellenauswahl'):
                    builder.sources(folder)

    def test_email_attachments_are_real_current_case_bytes(self):
        for c in CASES:
            directory=ROOT/'testakten'/c['slug']
            for item in documents(c):
                if not item['file'].endswith('.eml'):continue
                with self.subTest(case=c['slug'],mail=item['file']):
                    msg=BytesParser(policy=policy.default).parsebytes((directory/item['file']).read_bytes())
                    self.assertTrue(msg['From'] and msg['To'] and msg['Date'] and msg['Message-ID'])
                    actual={a.get_filename():a.get_payload(decode=True) for a in msg.iter_attachments()}
                    expected={Path(p).name:(directory/p).read_bytes() for p in item.get('attachments',[])}
                    self.assertEqual(actual,expected)
                    self.assertGreater(len(msg.get_body(preferencelist=('plain',)).get_content().strip()),300)

    def test_archives_preserve_five_stations_and_source_bytes(self):
        dist=Path(os.environ.get('BAUVERGABE_ZIPS','/tmp/bauvergabe-20261006/case-pdfs'))
        strict='BAUVERGABE_ZIPS' in os.environ
        for c in CASES:
            folder=ROOT/'testakten'/c['slug']
            self.assertTrue(preserves_directories(folder))
            self.assertTrue(set(STAGES)<= {p.name for p in folder.iterdir() if p.is_dir()})
            sources=working_dump_archive_pairs(folder,include_gesamt_pdf=True)
            self.assertEqual(len(sources),len({n for _,n in sources}))
            self.assertFalse(any(p.suffix in {'.md','.yaml'} for p,_ in sources))
            original=dist/f"testakte-{c['slug']}.zip"
            individual=dist/f"testakte-{c['slug']}-einzelpdfs.zip"
            if strict:
                self.assertTrue(original.is_file(),f'Native ZIP fehlt: {original}')
                self.assertTrue(individual.is_file(),f'Einzel-PDF-ZIP fehlt: {individual}')
                self.assertEqual(len(sources),NATIVE_COUNTS[c['slug']]+1)
                self.assertIn(folder/'gesamt-pdf'/f"{c['slug']}_gesamt.pdf",{p for p,_ in sources})
            if original.exists():
                with zipfile.ZipFile(original) as z:
                    self.assertEqual(len(z.namelist()),len(set(z.namelist())))
                    self.assertEqual(z.namelist()[0],'README.txt')
                    self.assertEqual(z.read('README.txt'),NOTICE_BYTES)
                    self.assertEqual(set(z.namelist()),{'README.txt',*(n for _,n in sources)})
                    for p,n in sources:self.assertEqual(z.read(n),p.read_bytes())
            if individual.exists():
                with zipfile.ZipFile(individual) as z:
                    self.assertEqual(len(z.namelist()),len(set(z.namelist())))
                    self.assertEqual(z.read('README.txt'),NOTICE_BYTES)
                    expected_pdfs=document_arcname_pairs(folder)
                    self.assertEqual(len(expected_pdfs),NATIVE_COUNTS[c['slug']])
                    self.assertEqual(set(z.namelist()),{'README.txt',*(name for _,name in expected_pdfs)})
                    self.assertTrue(set(STAGES)<={n.split('/')[0] for n in z.namelist() if n.endswith('.pdf')})
                    self.assertTrue(all(n=='README.txt' or n.endswith('.pdf') for n in z.namelist()))
                    self.assertIsNone(z.testzip())
                    for source,name in expected_pdfs:
                        self.assertTrue(z.read(name).startswith(b'%PDF-'))
                        if source.suffix=='.pdf':self.assertEqual(z.read(name),source.read_bytes())

if __name__=='__main__':unittest.main()
