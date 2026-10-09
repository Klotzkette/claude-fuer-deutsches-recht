#!/usr/bin/env python3
"""Native Artikel-5- und Anschlussakten aus derselben kontrollierten Quelle erzeugen.

XLSX zuvor mit gleichnamigem MJS-Builder; DOCX/PDF und echte MIME-Anhänge hier.
Keine zentralen Archive und keine rechtlichen Musterlösungen im Nutzercorpus.
"""
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.headerregistry import Address
from email.utils import getaddresses
from concurrent.futures import ThreadPoolExecutor
import json, hashlib, argparse, subprocess, os, sys, importlib.util, io, zipfile, re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
RUNTIME=Path(os.environ.get('CODEX_PRIMARY_RUNTIME','/Users/klotzkette/.cache/codex-runtimes/codex-primary-runtime'))
RENDER=Path('/Users/klotzkette/.codex/plugins/cache/openai-primary-runtime/documents/26.1007.11041/skills/documents/render_docx.py')
WARNING='> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n>\n> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'
TAG='ki-verordnung-v445.35.0'
def word(item,path):
 d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=sec.right_margin=Cm(2.2);sec.top_margin=sec.bottom_margin=Cm(1.8)
 for st in d.styles:
  if st.type==1:st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
 for name,size in [('Title',15),('Heading 1',11)]:
  st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_after=Pt(7);st.paragraph_format.space_before=Pt(7);st.paragraph_format.keep_with_next=True
 d.styles['Normal'].paragraph_format.line_spacing=1.05;d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.widow_control=True
 for n in list(d.styles.element.iter()):
  if n.tag==qn('w:pBdr'):n.getparent().remove(n)
  if n.tag==qn('w:rFonts'):
   for k in list(n.attrib):
    if 'theme' in k.lower():del n.attrib[k]
   n.set(qn('w:ascii'),'Times New Roman');n.set(qn('w:hAnsi'),'Times New Roman')
 d.add_paragraph(item['sender']);d.add_paragraph('An: '+item['recipient']);d.add_paragraph(item['date']);d.add_paragraph(item['title'],'Title')
 for block in re.split(r'\n\s*\n',item['body']):d.add_paragraph(block[3:] if block.startswith('## ') else block,'Heading 1' if block.startswith('## ') else 'Normal')
 p=sec.footer.paragraphs[0];p.alignment=2;p.add_run('Seite ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
 d.core_properties.title=item['title'];d.core_properties.author=item['sender'].splitlines()[0];d.core_properties.created=d.core_properties.modified=datetime(2026,10,9,12)
 b=io.BytesIO();d.save(b)
 with zipfile.ZipFile(io.BytesIO(b.getvalue())) as z,zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as out:
  for n in sorted(z.namelist()):i=zipfile.ZipInfo(n,(2026,10,9,12,0,0));i.compress_type=zipfile.ZIP_DEFLATED;out.writestr(i,z.read(n))
def eml(x,d):
 m=EmailMessage(policy=SMTP)
 for h,k in [('From','sender'),('To','recipient')]:m[h]=tuple(Address(display_name=n,addr_spec=a) for n,a in getaddresses([x[k]]))
 m['Subject']=x['title'];m['Date']=datetime.fromisoformat(x['date']);key=hashlib.sha256((d.name+x['file']).encode()).hexdigest()[:24];m['Message-ID']=f'<{key}@ki-akte.example>';m.set_content(x['body']+'\n',charset='utf-8')
 ty={'.docx':'vnd.openxmlformats-officedocument.wordprocessingml.document','.xlsx':'vnd.openxmlformats-officedocument.spreadsheetml.sheet','.pdf':'pdf'}
 for name in x['attachments']:p=d/name;m.add_attachment(p.read_bytes(),maintype='application',subtype=ty[p.suffix],filename=name)
 if x['attachments']:m.set_boundary('ki-akte-'+key)
 (d/x['file']).write_bytes(m.as_bytes())
def render(job,qa,env):
 item,p=job;out=qa/'word'/p.parent.name/p.stem;out.mkdir(parents=True,exist_ok=True)
 r=subprocess.run([sys.executable,str(RENDER),str(p),'--output_dir',str(out),'--emit_pdf','--dpi','100'],text=True,capture_output=True,env=env,timeout=300);(out/'render.log').write_text(r.stdout+r.stderr)
 png=sorted(out.glob('page-*.png'));pdf=list(out.glob('*.pdf'))
 if r.returncode or not png or len(pdf)!=1:raise RuntimeError(str(p)+' '+r.stderr[-1000:])
 p.with_suffix('.pdf').write_bytes(pdf[0].read_bytes());return {'file':str(p.relative_to(ROOT)),'pages':len(png),'images':[str(x) for x in png]}
def main():
 ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--qa-dir',type=Path,default=Path('/tmp/art5-qa'));ap.add_argument('--skip-render',action='store_true');ap.add_argument('--only',nargs='*');a=ap.parse_args();a.qa_dir.mkdir(parents=True,exist_ok=True)
 cases=json.loads((ROOT/'scripts/data/ki-verordnung-artikel5-testakten.json').read_text());cases=[c for c in cases if not a.only or c['slug'] in a.only]
 spec=importlib.util.spec_from_file_location('render_helpers',ROOT/'scripts/render-startup-gruender-werkstatt.py');helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper);env=helper.renderer_environment(a.qa_dir,RUNTIME);env['PYTHONPATH']=os.environ.get('PYTHONPATH','')
 jobs=[]
 for c in cases:
  d=ROOT/'testakten'/c['slug'];d.mkdir(parents=True,exist_ok=True)
  for item in c['docs']:p=d/item['file'];word(item,p);jobs.append((item,p))
  (d/'12_Projektchat.txt').write_text(c['chat'])
 if not a.skip_render:
  with ThreadPoolExecutor(max_workers=2) as pool:reports=list(pool.map(lambda j:render(j,a.qa_dir,env),jobs))
  old=json.loads((a.qa_dir/'word-render.json').read_text()) if (a.qa_dir/'word-render.json').exists() else [];slugs={c['slug'] for c in cases};reports=[x for x in old if x['file'].split('/')[1] not in slugs]+reports;(a.qa_dir/'word-render.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2)+'\n')
 manifest=[]
 for c in cases:
  d=ROOT/'testakten'/c['slug']
  for x in c['emails']:eml(x,d)
  files=sorted(p for p in d.iterdir() if p.is_file() and p.suffix.lower() in ['.docx','.pdf','.eml','.xlsx','.txt'] and p.name!='README.txt')
  base=f'https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/{TAG}'
  readme=f"# Testakte {c['title']}\n\n## 1. Aufgabe\n\n{c['brief']}\n\nAlle Personen, Organisationen, Anschriften, Vorgänge und Daten sind fiktiv. Aktenstand ist der 09.10.2026. Zukunftsangaben sind Planungen oder ausdrücklich bezeichnete Übungsszenarien. Der Nutzercorpus enthält keine rechtliche Musterlösung.\n\n## 2. Bearbeitung\n\nBeginnen Sie mit dem Prüfauftrag. Lesen Sie die technischen beziehungsweise organisatorischen Angaben, die E-Mails mit ihren echten Anhängen und die Rechenmappe. Erstellen Sie das bestellte ausformulierte Produkt; halten Sie widersprüchliche Aussagen und entscheidende fehlende Tatsachen offen. Excel-Ergebnisse sind Rechenbelege und keine automatische rechtliche Entscheidung.\n\n## 3. Downloads\n\n{WARNING}\n\n| Variante | Download |\n|---|---|\n| Originaldateien | [Originalformate als ZIP]({base}/{c['slug']}-original.zip) |\n| Einzelne Lesefassungen | [Einzel-PDFs als ZIP]({base}/{c['slug']}-einzel-pdf.zip) |\n| Gesamtakte | [Gesamt-PDF]({base}/{c['slug']}_gesamt.pdf) |\n\n## 4. Bestand\n\n{WARNING}\n\n| Datei | Format |\n|---|---|\n"+'\n'.join(f'| [{p.name}]({p.name}) | {p.suffix[1:].upper()} |' for p in files)+'\n\nWord-Dateien und zugehörige PDF-Fassungen haben denselben Inhalt. Anhänge sind zusätzlich in den E-Mails enthalten; sie sind keine weiteren unabhängigen Beweise.\n'
  (d/'README.md').write_text(readme)
  manifest.append({'slug':c['slug'],'title':c['title'],'brief':c['brief'],'files':[{'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size} for p in files]})
 out=ROOT/'quality/ki-verordnung-2026-10-09/artikel5/aktenmanifest.json';out.parent.mkdir(parents=True,exist_ok=True)
 old=json.loads(out.read_text()) if out.exists() else [];slugs={c['slug'] for c in cases};out.write_text(json.dumps([x for x in old if x['slug'] not in slugs]+manifest,ensure_ascii=False,indent=2)+'\n')
 print(len(cases),'Akten;',sum(len(x['files']) for x in manifest),'Dateien')
if __name__=='__main__':main()
