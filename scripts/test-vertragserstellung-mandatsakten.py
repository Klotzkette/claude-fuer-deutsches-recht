#!/usr/bin/env python3
"""Prueft Originalunterlagen, Anlagen und Zahlen der vier Vertragsmandate."""
from collections import Counter
from datetime import datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path
import re
import unittest

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from testakte_file_filter import include_in_working_dump
from testakte_zip_common import working_dump_archive_pairs

ROOT = Path(__file__).resolve().parent.parent
SLUGS = (
    'vertragserstellung-gbr-fahrradwerkstatt-leipzig',
    'vertragserstellung-ladenkooperation-keramik-freiburg',
    'vertragserstellung-softwarelizenz-pumpen-bremen',
    'vertragserstellung-ug-eventkueche-regensburg',
)


class ContractFiles(unittest.TestCase):
    def test_inventory_and_exports(self):
        for slug in SLUGS:
            with self.subTest(slug=slug):
                case = ROOT / 'testakten' / slug
                sources = sorted(p for p in case.iterdir() if p.is_file() and p.name[:2].isdigit())
                self.assertEqual([p.name[:2] for p in sources], [f'{i:02}' for i in range(1, 13)])
                self.assertEqual(Counter(p.suffix for p in sources), {'.docx': 4, '.eml': 4, '.txt': 3, '.xlsx': 1})
                readme = (case / 'README.md').read_text()
                for path in sources:
                    self.assertIn(f']({path.name})', readme)
                self.assertFalse(include_in_working_dump(case / 'rubric.yaml', case))
                exported = [name for _, name in working_dump_archive_pairs(case, include_gesamt_pdf=True)]
                self.assertEqual(len(exported), 13)
                self.assertTrue(all('/' not in n and not n.endswith(('.md', '.yaml')) for n in exported))

    def test_emails_and_embedded_attachments(self):
        ids = set()
        for slug in SLUGS:
            case = ROOT / 'testakten' / slug
            for path in case.glob('*.eml'):
                with self.subTest(path=path.name):
                    mail = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                    for header in ('From', 'To', 'Date', 'Subject', 'Message-ID', 'MIME-Version'):
                        self.assertTrue(mail[header], header)
                    self.assertNotIn(mail['Message-ID'], ids)
                    ids.add(mail['Message-ID'])
                    body = mail.get_body(preferencelist=('plain',)).get_content()
                    self.assertGreater(len(body), 600)
                    self.assertRegex(body, '[äöüÄÖÜß]')
                    self.assertFalse(mail.defects)
                    for part in mail.iter_attachments():
                        self.assertEqual(part.get_payload(decode=True), (case / part.get_filename()).read_bytes())

    def test_documents_are_complete_sources(self):
        for slug in SLUGS:
            for path in (ROOT / 'testakten' / slug).glob('*.docx'):
                with self.subTest(path=path.name):
                    doc = Document(path)
                    text = '\n'.join(p.text for p in doc.paragraphs)
                    self.assertGreater(len(text), 900)
                    self.assertRegex(text, '[äöüÄÖÜß]')
                    self.assertNotRegex(text, r'(?i)Testakte|fiktiv|Platzhalter|Formathinweis|Lösungsmatrix')
                    self.assertNotIn(chr(167), text)
                    self.assertEqual(doc.styles['Normal'].font.name, 'Times New Roman')
                    self.assertEqual(doc.styles['Normal'].font.size.pt, 11)
                    for para in doc.paragraphs:
                        if para.style.name.startswith('Heading'):
                            self.assertRegex(para.text, r'^\d+\. ')

    def test_spreadsheet_sources_and_saved_calculations(self):
        for slug in SLUGS:
            path = next((ROOT / 'testakten' / slug).glob('*.xlsx'))
            book = load_workbook(path, data_only=True)
            self.assertEqual(len(book.sheetnames), 2)
            for sheet in book:
                self.assertNotRegex(' '.join(str(c.value) for row in sheet for c in row), r'#REF!|#DIV/0!|#VALUE!')
            if 'gbr-' in slug:
                self.assertEqual(book['Einzahlungen']['B10'].value, 18000)
                self.assertEqual(book['Einzahlungen']['C10'].value, 0)
            elif '-ug-' in slug:
                self.assertEqual(book['Kapital']['B9'].value, 1000)
                self.assertEqual(book['Kapital']['C9'].value, 0)
            elif 'softwarelizenz' in slug:
                self.assertEqual(sum(book['Benutzer'].cell(r, 3).value for r in range(6, 10)), 24)
                self.assertEqual(book['Benutzer']['C10'].value, 4)
                self.assertEqual(sum(book['Entgelte'].cell(r, 4).value for r in (6, 7)), 27132)
            else:
                self.assertEqual(sum(book['Lieferbestand'].cell(r, 3).value for r in range(6, 14)), 41)
                self.assertAlmostEqual(book['Kassenjournal']['E11'].value, 119)
                self.assertIsInstance(book['Kassenjournal']['B6'].value, datetime)
            book.close()

    def test_combined_pdfs_are_complete_and_unmarked(self):
        for slug in SLUGS:
            with self.subTest(slug=slug):
                pdf = PdfReader(ROOT / 'testakten' / slug / 'gesamt-pdf' / f'{slug}_gesamt.pdf')
                text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
                self.assertGreaterEqual(len(pdf.pages), 14)
                self.assertGreater(len(text), 18000)
                self.assertNotIn('Diese Testakte wurde', text)
                self.assertNotIn('human_review', text)
                self.assertNotIn('Lösungsmatrix', text)
                self.assertRegex(text, '[äöüÄÖÜß]')


if __name__ == '__main__':
    unittest.main()
