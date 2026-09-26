#!/usr/bin/env python3
"""Native Nachrechnung und Sichtprüfungsunterlagen für die Akten 4 bis 6. Autor: Klotzkette."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from akten_build_runtime import screen_font
from office_process import run_office
from testakte_office_pdf import office_binary, render_office_batch

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
CORE = "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"
DC = "{http://purl.org/dc/elements/1.1/}"
SLUGS = {4: "bauwirtschaft-hoai-4-genehmigung-werkhof-celle", 5: "bauwirtschaft-hoai-5-ausfuehrung-schule-hameln", 6: "bauwirtschaft-hoai-6-lv-sporthalle-peine"}
# Stabile Erwartungen stammen aus den geometrischen Ansätzen der Originalunterlagen.
SPECS = [
    (4, "06_Flaechen.xlsx", [("Flaechen", "A1:F29")], {"Flaechen!E13": 1438, "Flaechen!E14": 1442, "Flaechen!E20": 448.8}, ("Flaechen!D18", 14.6, {"Flaechen!E20": 446.4}), ("Flaechen!D18", None, {"Flaechen!E20": None})),
    (5, "08_Hoehenketten.xlsx", [("Hoehen", "A1:H25")], {"Hoehen!G6": 2.68, "Hoehen!G7": 3.13, "Hoehen!H6": -.07}, ("Hoehen!B6", 3.25, {"Hoehen!G6": 2.78}), ("Hoehen!D6", None, {"Hoehen!G6": None})),
    (6, "03_Kostenberechnung.xlsx", [("Kosten", "A1:C23")], {"Kosten!B13": 1065000, "Kosten!B17": 1344700}, ("Kosten!B6", 46000, {"Kosten!B17": 1345890}), ("Kosten!B14", 0, {"Kosten!B17": 1267350})),
    (6, "06_Mengenermittlung.xlsx", [("Mengen", "A1:H27")], {"Mengen!G8": 138.96, "Mengen!G9": 286.56, "Mengen!G10": 100.1746, "Mengen!G11": 200.3492, "Mengen!G13": 396}, ("Mengen!B16", 6, {"Mengen!G9": 285.84}), ("Mengen!D6", None, {"Mengen!G6": None, "Mengen!G13": None})),
    (6, "08_Planer_LV.xlsx", [("LV", "A1:F27"), ("Kosten", "A1:C25")], {"LV!F23": 92282.86, "Kosten!B8": -1717.14, "Kosten!B15": 1128282.86}, ("LV!E6", 13, {"LV!F23": 92570.86, "Kosten!B15": 1128570.86}), ("LV!E6", None, {"LV!F23": None, "Kosten!B15": None})),
    (6, "13_Vergabetermine.xlsx", [("Termine", "A1:E24")], {"Termine!C15": 10}, ("Termine!B8", 42, {"Termine!C15": 3}), ("Termine!B8", 0, {"Termine!C15": 45})),
]


def module(name, filename):
    spec = importlib.util.spec_from_file_location(name, HERE / filename)
    result = importlib.util.module_from_spec(spec)
    sys.modules[name] = result
    spec.loader.exec_module(result)
    return result


def original_files(phase):
    return sorted(p for p in (ROOT / "testakten" / SLUGS[phase]).iterdir() if p.is_file() and re.fullmatch(r"\d{2}_[A-Za-z0-9_.-]+", p.name))


def patch_workbook(path, areas, mutation=None, clear=False):
    """Nur Metadaten/Drucklayout bzw. eine Testeingabe; keine Formelautorenschaft."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item)
            if item.filename == "docProps/core.xml":
                root = ET.fromstring(data)
                for tag in (DC + "creator", CORE + "lastModifiedBy"):
                    el = root.find(tag)
                    if el is None: el = ET.SubElement(root, tag)
                    el.text = "Klotzkette"
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif item.filename == "xl/workbook.xml":
                root = ET.fromstring(data)
                names = root.find(NS + "definedNames")
                if names is None: names = ET.SubElement(root, NS + "definedNames")
                for el in list(names):
                    if el.get("name") in ("_xlnm.Print_Area", "_xlnm.Print_Titles"): names.remove(el)
                for index, (sheet, area) in enumerate(areas):
                    ET.SubElement(names, NS + "definedName", name="_xlnm.Print_Area", localSheetId=str(index)).text = f"'{sheet}'!{area}"
                    if sheet == "LV":
                        ET.SubElement(names, NS + "definedName", name="_xlnm.Print_Titles", localSheetId=str(index)).text = "'LV'!$1:$5"
                calc = root.find(NS + "calcPr")
                if calc is None: calc = ET.SubElement(root, NS + "calcPr")
                calc.attrib.update(calcMode="auto", fullCalcOnLoad="1", forceFullCalc="1")
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif re.fullmatch(r"xl/worksheets/sheet\d+\.xml", item.filename):
                index = int(re.search(r"sheet(\d+)", item.filename)[1]) - 1
                root = ET.fromstring(data)
                setup = root.find(NS + "pageSetup")
                if setup is None: setup = ET.SubElement(root, NS + "pageSetup")
                setup.attrib.update(paperSize="8", orientation="landscape", fitToWidth="1", fitToHeight="0" if areas[index][0] == "LV" else "1")
                margins = root.find(NS + "pageMargins")
                if margins is None: margins = ET.SubElement(root, NS + "pageMargins")
                margins.attrib.update(left="0.3", right="0.3", top="0.35", bottom="0.35", header="0.1", footer="0.1")
                prop = root.find(NS + "sheetPr")
                if prop is None: prop = ET.Element(NS + "sheetPr"); root.insert(0, prop)
                pg = prop.find(NS + "pageSetUpPr")
                if pg is None: pg = ET.SubElement(prop, NS + "pageSetUpPr")
                pg.set("fitToPage", "1")
                if mutation and mutation[0].split("!")[0] == areas[index][0]:
                    cell = root.find(f".//{NS}c[@r='{mutation[0].split('!')[1]}']")
                    if cell is None: raise ValueError(mutation)
                    for child in list(cell): cell.remove(child)
                    cell.attrib.pop("t", None)
                    if mutation[1] is not None: ET.SubElement(cell, NS + "v").text = str(mutation[1])
                if clear or mutation:
                    for cell in root.findall(f".//{NS}c"):
                        if cell.find(NS + "f") is not None:
                            for el in cell.findall(NS + "v"): cell.remove(el)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            dst.writestr(item, data)
    path.write_bytes(buffer.getvalue())


def formulas(path):
    wb = load_workbook(path, data_only=False)
    result = {f"{s.title}!{c.coordinate}": c.value for s in wb for row in s for c in row if c.data_type == "f"}
    wb.close()
    return result


def check_values(path, expected):
    wb = load_workbook(path, data_only=True)
    observed = {}
    for address, target in expected.items():
        sheet, cell = address.split("!"); value = wb[sheet][cell].value
        assert (value in (None, "") if target is None else isinstance(value, (int, float)) and abs(value - target) < .000001), (path.name, address, value, target)
        observed[address] = value
    assert not [(s.title, c.coordinate, c.value) for s in wb for row in s for c in row if c.data_type == "e"], path
    wb.close()
    return observed


def roundtrip(paths, folder):
    output = folder / "output"; output.mkdir(parents=True, exist_ok=True)
    binary = office_binary()
    if not binary: raise RuntimeError("SOFFICE oder LibreOffice im PATH erforderlich")
    command = [binary, f"-env:UserInstallation={(folder / 'profile').resolve().as_uri()}", "--headless", "--convert-to", "xlsx", "--outdir", str(output), *map(str, paths)]
    code, log = run_office(command, env=os.environ.copy(), timeout=240)
    (folder / "office.log").write_text(log, encoding="utf-8")
    result = [output / p.name for p in paths]
    if code or not all(p.is_file() for p in result): raise RuntimeError(log)
    return result


def recalculate(qa, phases):
    specs = [spec for spec in SPECS if spec[0] in phases]
    paths = [ROOT / "testakten" / SLUGS[spec[0]] / spec[1] for spec in specs]
    original = []
    for source, spec in zip(paths, specs):
        patch_workbook(source, spec[2], clear=True); original.append(formulas(source))
    results = roundtrip(paths, qa / "native-baseline")
    evidence = []
    for source, result, spec, before in zip(paths, results, specs, original):
        assert formulas(result) == before, (source.name, "Formeln verändert")
        values = check_values(result, spec[3])
        shutil.copyfile(result, source); patch_workbook(source, spec[2])
        evidence.append(dict(phase=spec[0], file=source.name, formula_count=len(before), baseline=values, mutations=[]))
    for test_index, label in ((4, "number"), (5, "boundary")):
        inputs = []
        for source, spec in zip(paths, specs):
            target = qa / ("mutation-" + label) / "input" / source.name
            target.parent.mkdir(parents=True, exist_ok=True); shutil.copyfile(source, target)
            patch_workbook(target, spec[2], mutation=spec[test_index]); inputs.append(target)
        for result, spec, before, record in zip(roundtrip(inputs, qa / ("mutation-" + label)), specs, original, evidence):
            assert formulas(result) == before, result
            mutation = spec[test_index]
            record["mutations"].append(dict(input_cell=mutation[0], input_value=mutation[1], observed=check_values(result, mutation[2])))
    for record, source in zip(evidence, paths): record["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    (qa / "native-recalculation.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(evidence)} native Basisrechnungen und {len(evidence) * 2} Eingabemutationen bestanden", flush=True)


def normalize_pdf(data, title):
    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(data)).pages: writer.add_page(page)
    writer.add_metadata({"/Author": "Klotzkette", "/Creator": "Klotzkette", "/Producer": "Klotzkette", "/Title": title})
    out = io.BytesIO(); writer.write(out); return out.getvalue()


def render(qa, phases):
    exporter = module("hoai456_pdf", "build-testakten-einzelpdf-zips.py")
    renderer = os.environ.get("DOCX_RENDERER")
    if not renderer or not Path(renderer).is_file(): raise RuntimeError("DOCX_RENDERER auf render_docx.py der Dokumentwerkzeuge setzen")
    poppler = shutil.which("pdftoppm")
    if not poppler: raise RuntimeError("Poppler pdftoppm im PATH erforderlich")
    files = [(phase, path) for phase in phases for path in original_files(phase)]
    office = render_office_batch([path for _, path in files if path.suffix == ".xlsx"])
    index = []
    for phase, source in files:
        dest = qa / f"phase-{phase}" / source.stem; dest.mkdir(parents=True, exist_ok=True)
        if source.suffix == ".docx":
            for old in (dest / "docx").glob("page-*.png"): old.unlink()
            command = [os.environ.get("DOCX_PYTHON", sys.executable), renderer, str(source), "--output_dir", str(dest / "docx"), "--emit_pdf", "--dpi", "110"]
            result = subprocess.run(command, capture_output=True, text=True)
            (dest / "renderer.log").write_text(result.stdout + result.stderr, encoding="utf-8")
            if result.returncode: raise RuntimeError(result.stdout + result.stderr)
            data = (dest / "docx" / (source.stem + ".pdf")).read_bytes()
        else: data = exporter.render_document_pdf(source, source.parent, office)
        if not data: raise RuntimeError(f"Leeres Render-Ergebnis: {source}")
        pdf = dest / (source.stem + ".pdf"); pdf.write_bytes(normalize_pdf(data, source.stem))
        for old in dest.glob("page-*.png"): old.unlink()
        subprocess.run([poppler, "-png", "-r", "110", str(pdf), str(dest / "page")], check=True, capture_output=True)
        images = sorted(dest.glob("page-*.png")); pages = len(PdfReader(pdf).pages)
        assert pages == len(images)
        index.append(dict(phase=phase, file=source.name, sha256=hashlib.sha256(source.read_bytes()).hexdigest(), pages=pages, pdf=str(pdf), images=list(map(str, images))))
        print(f"LPH {phase} / {source.name}: {pages} Seiten", flush=True)
    (qa / "render-index.json").write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    for phase in phases:
        images = [(item["file"], p) for item in index if item["phase"] == phase for p in item["images"]]
        for offset in range(0, len(images), 6):
            contact = Image.new("RGB", (1800, 3 * 1230), "#d7d7d7"); draw = ImageDraw.Draw(contact)
            for number, (name, image) in enumerate(images[offset:offset + 6]):
                im = Image.open(image).convert("RGB"); im.thumbnail((875, 1170))
                x = number % 2 * 900; y = number // 2 * 1230
                contact.paste(im, (x + (900 - im.width) // 2, y + 40))
                draw.text((x + 15, y + 9), f"{name} / {Path(image).stem}", font=screen_font(20), fill="black")
            contact.save(qa / f"contact-{phase}-{offset // 6 + 1}.jpg", quality=92)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--qa-dir", type=Path, default=Path(tempfile.gettempdir()) / "hoai-4-6-qa")
    parser.add_argument("--phase", type=int, choices=[4, 5, 6], action="append")
    parser.add_argument("--stage", choices=["all", "recalculate", "render"], default="all")
    args = parser.parse_args(); args.qa_dir.mkdir(parents=True, exist_ok=True)
    phases = args.phase or [4, 5, 6]
    if args.stage in ("all", "recalculate"): recalculate(args.qa_dir, phases)
    if args.stage in ("all", "render"): render(args.qa_dir, phases)
    print(f"QA-Ausgabe: {args.qa_dir}")


if __name__ == "__main__": main()
