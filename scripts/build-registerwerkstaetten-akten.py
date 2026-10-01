#!/usr/bin/env python3
"""Baut die nativen Unterlagen der vier Registerwerkstätten aus lesbaren Falldaten.

Die XLSX-Spezifikation wird für build-registerakten-tabellen.mjs ausgegeben.
Gesamt-PDF und beide ZIP-Varianten entstehen mit den bestehenden Repo-Buildern.
"""
from __future__ import annotations
import argparse
import csv
from datetime import datetime, timezone
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
import importlib
import io
import json
from pathlib import Path
import re
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
MODULES = {'transparenz': 'registerakten_transparenz', 'handelsregister': 'registerakten_handelsregister', 'grundbuch': 'registerakten_grundbuch', 'marken': 'registerakten_marken'}
FIXED = datetime(2026, 10, 1, 9, 0, tzinfo=timezone.utc)
FONTDIR = Path('/System/Library/Fonts/Supplemental')


def paragraphs(item):
    body = item.get('body', '')
    return body if isinstance(body, list) else re.split(r'\n\s*\n', body.strip())


def stable_docx(doc, target):
    doc.core_properties.created = doc.core_properties.modified = FIXED
    memory = io.BytesIO(); doc.save(memory)
    with zipfile.ZipFile(memory) as source, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2026, 10, 1, 9, 0, 0)); info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, source.read(name))


def word(item, target):
    doc = Document(); section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.1)
    section.left_margin = section.right_margin = Cm(2.3)
    for name, size in [('Normal', 11), ('Title', 17), ('Heading 1', 13), ('Heading 2', 12)]:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_after = Pt(7); st.paragraph_format.line_spacing = 1.12
        st.paragraph_format.widow_control = True
        if name != 'Normal': st.paragraph_format.keep_with_next = True
    doc.core_properties.title = item['title']; doc.core_properties.author = item.get('sender', 'Aktenverwaltung')
    if item.get('sender'): doc.add_paragraph(item['sender'])
    if item.get('recipient'): doc.add_paragraph(item['recipient'])
    doc.add_paragraph(item['title'], 'Title')
    metadata = ' · '.join(str(item[k]) for k in ('date', 'reference') if item.get(k))
    if metadata: doc.add_paragraph(metadata)
    for text in paragraphs(item):
        if text.startswith('## '): doc.add_paragraph(text[3:], 'Heading 1')
        else: doc.add_paragraph(text)
    for spec in item.get('tables', []):
        table = doc.add_table(rows=1, cols=len(spec['headers'])); table.autofit = False
        table.style = 'Table Grid'
        for cell, text in zip(table.rows[0].cells, spec['headers']): cell.text = str(text)
        for values in spec['rows']:
            cells = table.add_row().cells
            for cell, value in zip(cells, values): cell.text = str(value)
        props = table.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); props.append(repeat)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    for run in p.runs: run.font.size = Pt(10)
    footer = section.footer.paragraphs[0]; footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    stable_docx(doc, target)


def pdf(item, target):
    if 'RegisterTimes' not in pdfmetrics.getRegisteredFontNames():
        if (FONTDIR / 'Times New Roman.ttf').exists():
            pdfmetrics.registerFont(TTFont('RegisterTimes', str(FONTDIR / 'Times New Roman.ttf')))
            pdfmetrics.registerFont(TTFont('RegisterTimesBold', str(FONTDIR / 'Times New Roman Bold.ttf')))
        else:
            # Standard PDF Times fonts cover all German characters used here.
            pass
    family = 'RegisterTimes' if 'RegisterTimes' in pdfmetrics.getRegisteredFontNames() else 'Times-Roman'
    bold = 'RegisterTimesBold' if family == 'RegisterTimes' else 'Times-Bold'
    body = ParagraphStyle('Body', fontName=family, fontSize=11, leading=14, spaceAfter=8)
    title = ParagraphStyle('Title', parent=body, fontName=bold, fontSize=16, leading=20, spaceAfter=13)
    small = ParagraphStyle('Small', parent=body, fontSize=9, leading=11)
    wide = any(len(t['headers']) > 5 for t in item.get('tables', []))
    page = landscape(A4) if wide else A4; width = page[0] - 4.4*cm
    def para(text, style=body): return Paragraph(escape(str(text)).replace('\n', '<br/>'), style)
    story = []
    if item.get('sender'): story.append(para(item['sender']))
    if item.get('recipient'): story.append(para(item['recipient']))
    story += [para(item['title'], title)]
    meta = ' · '.join(str(item[k]) for k in ('date', 'reference') if item.get(k))
    if meta: story.append(para(meta, small))
    story.append(Spacer(1, 6))
    for p in paragraphs(item): story.append(para(p[3:] if p.startswith('## ') else p, title if p.startswith('## ') else body))
    for spec in item.get('tables', []):
        rows = [[para(v, small) for v in spec['headers']]] + [[para(v, small) for v in row] for row in spec['rows']]
        table = Table(rows, colWidths=[width/len(spec['headers'])]*len(spec['headers']), repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e8edf1')), ('GRID',(0,0),(-1,-1),.35,colors.HexColor('#c7cdd1')), ('VALIGN',(0,0),(-1,-1),'TOP'), ('TOPPADDING',(0,0),(-1,-1),6), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        story.extend([Spacer(1, 6), table, Spacer(1, 9)])
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont(family, 9); canvas.drawRightString(page[0]-2.2*cm, 1.2*cm, str(doc.page)); canvas.restoreState()
    SimpleDocTemplate(str(target), pagesize=page, leftMargin=2.2*cm, rightMargin=2.2*cm, topMargin=2.0*cm, bottomMargin=2.0*cm, title=item['title'], author=item.get('sender','Aktenverwaltung'), invariant=1).build(story, onFirstPage=footer, onLaterPages=footer)


def mail(item, target):
    msg = EmailMessage(policy=SMTP)
    msg['From'] = item['from']; msg['To'] = item['to']
    if item.get('cc'): msg['Cc'] = item['cc']
    msg['Subject'] = item.get('subject', item['title'])
    raw = item.get('date', '2026-10-01T09:00:00+02:00')
    try: stamp = datetime.fromisoformat(raw)
    except ValueError: stamp = datetime.strptime(raw, '%d.%m.%Y').replace(hour=9, tzinfo=timezone.utc)
    if stamp.tzinfo is None: stamp = stamp.replace(tzinfo=timezone.utc)
    msg['Date'] = format_datetime(stamp)
    msg['Message-ID'] = '<' + target.stem + '.' + target.parent.name + '@akten.example>'
    msg.set_content('\n\n'.join(paragraphs(item)), charset='utf-8'); target.write_bytes(msg.as_bytes())


def screenshot(item, target):
    # Internal working view, no government portal branding, seal or signature.
    fontpath = FONTDIR / 'Arial.ttf'
    normal = ImageFont.truetype(str(fontpath), 24) if fontpath.exists() else ImageFont.load_default(size=24)
    heading = ImageFont.truetype(str(fontpath), 30) if fontpath.exists() else normal
    texts = [item['title'], item.get('date','')] + paragraphs(item)
    if item.get('headers'): texts += [' | '.join(map(str,item['headers']))] + [' | '.join(map(str,r)) for r in item.get('rows',[])]
    scratch = Image.new('RGB',(1440,100)); draw = ImageDraw.Draw(scratch); lines=[]
    for text in texts:
        for para in str(text).splitlines():
            words = para.split(); line = ''
            for word in words:
                nxt = (line+' '+word).strip()
                if draw.textlength(nxt,font=normal)>1280 and line: lines.append(line); line=word
                else: line=nxt
            lines.append(line)
        lines.append('')
    image = Image.new('RGB',(1440,max(700,180+len(lines)*36)),'#f4f5f7'); draw=ImageDraw.Draw(image)
    draw.rectangle((0,0,1440,70),fill='#263641'); draw.text((54,19),'Arbeitsablage  |  Verlauf',fill='white',font=heading)
    y=116
    for line in lines: draw.text((60,y),line,fill='#172630',font=normal);y+=36
    image.save(target)


def readme(case, directory):
    lines=[f"# 1. {case['title']}",'', '## 1.1. Vorgang','',case['summary'],'',f"Stand: {case['date']}. Mandat: {case['client']}.",'',f"Passendes Plugin: [`{case['plugin']}`](../../{case['plugin']}/README.md).",'', '## 1.2. Arbeitsauftrag','',case['assignment'],'', '<!-- reserved-example-contacts -->','', 'Alle Personen und Unternehmen in dieser Akte sind fiktiv. Kontaktadressen mit `.example` sind nicht zustellbar. Die Unterlagen enthalten keine Musterlösung.','', '## 1.3. Unterlagen','', '| Datei | Inhalt |','| --- | --- |']
    lines += [f"| [{d['file']}]({d['file']}) | {d['title'].replace('|','/')} |" for d in case['documents']]
    directory.joinpath('README.md').write_text('\n'.join(lines)+'\n')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('groups', nargs='*', choices=list(MODULES))
    parser.add_argument('--sheet-spec', type=Path, required=True)
    args=parser.parse_args(); specs=[]; cases=[]
    for group in args.groups or MODULES:
        cases += importlib.import_module(MODULES[group]).CASES
    for case in cases:
        directory=ROOT/'testakten'/case['slug']; directory.mkdir(parents=True, exist_ok=True)
        names=[x['file'] for x in case['documents']]
        if len(names)!=len(set(names)) or any(Path(x).name!=x for x in names):raise ValueError('Dateikollision/Unterordner: '+case['slug'])
        for item in case['documents']:
            target=directory/item['file']; suffix=target.suffix.lower()
            if suffix=='.docx':word(item,target)
            elif suffix=='.pdf':pdf(item,target)
            elif suffix=='.eml':mail(item,target)
            elif suffix=='.png':screenshot(item,target)
            elif suffix=='.xlsx':specs.append({**item,'output':str(target),'case':case['slug']})
            elif suffix=='.txt':target.write_text('\n'.join(str(item[k]) for k in ['title','date','sender','recipient'] if item.get(k))+'\n\n'+'\n\n'.join(paragraphs(item))+'\n')
            elif suffix=='.csv':
                with target.open('w',newline='',encoding='utf-8-sig') as stream:
                    writer=csv.writer(stream,delimiter=';');writer.writerow(item['headers']);writer.writerows(item['rows'])
            else:raise ValueError('Unbekanntes Format: '+str(target))
        readme(case,directory);print(case['slug'],len(case['documents']),flush=True)
    args.sheet_spec.parent.mkdir(parents=True,exist_ok=True);args.sheet_spec.write_text(json.dumps(specs,ensure_ascii=False,indent=2)+'\n')


if __name__=='__main__':main()
