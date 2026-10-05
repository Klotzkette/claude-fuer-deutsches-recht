#!/usr/bin/env python3
"""Bestands- und Exportregression, keine automatische juristische Falllösung."""
import csv
from datetime import datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
import io
import os
from pathlib import Path
import re
import unittest
import zipfile

from docx import Document
from pypdf import PdfReader
import yaml

from verletzungsrisiko_akte_daten import SLUG, TITLE, DOCUMENTS, TEXTS, CSV_NAME, ATTACHMENTS
from testakte_disclaimer import NOTICE_BYTES

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT/'testakten'/SLUG
NAMES = {d['file'] for d in DOCUMENTS} | set(TEXTS) | {CSV_NAME}
PDF = CASE/'gesamt-pdf'/f'{SLUG}_gesamt.pdf'
PACKAGES = Path(os.environ.get('VERLETZUNG_PAKETE', ROOT/'dist'))


class Verletzungsrisiko(unittest.TestCase):
    def test_exact_original_inventory(self):
        actual = {p.name for p in CASE.iterdir() if p.is_file()} - {'README.md', 'rubric.yaml'}
        self.assertEqual(actual, NAMES)
        self.assertEqual(len(actual), 18)
        self.assertEqual(sum(n.endswith('.docx') for n in actual), 10)

    def test_word_documents_are_complete(self):
        for path in CASE.glob('*.docx'):
            doc = Document(path)
            text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertGreater(len(text), 900, path.name)
            self.assertTrue(any(p.style.name == 'Title' for p in doc.paragraphs), path.name)
            self.assertEqual(doc.styles['Normal'].font.name, 'Times New Roman')
            self.assertEqual(doc.styles['Normal'].font.size.pt, 11)

    def test_mail_headers_and_attachments(self):
        for path in CASE.glob('*.eml'):
            msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            self.assertFalse(msg.defects, path.name)
            for name in ('From', 'To', 'Date', 'Subject', 'Message-ID', 'MIME-Version'):
                self.assertIsNotNone(msg[name], (path.name, name))
                self.assertFalse(msg[name].defects, (path.name, name))
            self.assertLessEqual(parsedate_to_datetime(msg['Date']).date(), datetime(2026, 10, 5).date())
            self.assertGreater(len(msg.get_body(preferencelist=('plain',)).get_content()), 500)
            attached = {p.get_filename(): p.get_payload(decode=True) for p in msg.iter_attachments()}
            self.assertEqual(set(attached), set(ATTACHMENTS.get(path.name, [])))
            for name, data in attached.items():
                self.assertEqual(data, (CASE/name).read_bytes(), name)

    def test_costs_per_person(self):
        with (CASE/CSV_NAME).open(encoding='utf-8', newline='') as stream:
            reader = csv.DictReader(stream, delimiter=';')
            rows = list(reader)
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(None not in row and len(row) == 6 for row in rows))
        amount = lambda row: Decimal(row['Betrag_EUR'].replace(',', '.'))
        self.assertEqual(sum(map(amount, rows)), Decimal('403.20'))
        self.assertEqual(sum(amount(r) for r in rows if r['Person'] == 'Wolkenbein'), Decimal('249.00'))
        self.assertEqual(sum(amount(r) for r in rows if r['Person'] == 'Pimpinella'), Decimal('154.20'))
        for row in rows:
            self.assertLessEqual(datetime.strptime(row['Leistungstag'], '%d.%m.%Y'), datetime.strptime(row['Buchungstag'], '%d.%m.%Y'))

    def test_chronology_and_separate_receipts(self):
        data = {d['file']: d for d in DOCUMENTS}
        self.assertEqual(data['03_Fuhrpark.docx']['date'], '30.09.2026')
        self.assertTrue(data['14_Vertragsbestaetigung.eml']['date'].startswith('2026-09-29'))
        self.assertIn('wartet die Patientin', data['05_Ambulanz_Pimpinella.docx']['body'])
        self.assertTrue({'12a_Apotheke.docx', '12b_Massage.docx', '12c_Taxi.docx'}.issubset(NAMES))
        self.assertEqual(Decimal('600') + Decimal('31.20') + Decimal('5') + Decimal('79'), Decimal('715.20'))

    def test_originals_do_not_include_answers_or_training_notice(self):
        text = '\n'.join(d['body'] for d in DOCUMENTS) + '\n'.join(TEXTS.values())
        for forbidden in ('Musterlösung', 'Lösungsmatrix', 'Testakte', 'fiktiv', 'Platzhalter', 'Formathinweis', chr(167)):
            self.assertNotIn(forbidden.casefold(), text.casefold())
        self.assertRegex(text, r'[äöüÄÖÜß]')
        self.assertNotRegex(text, r'VI ZR \d')

    def test_readme_and_rubric(self):
        text = (CASE/'README.md').read_text()
        self.assertIn(TITLE, text)
        for name in NAMES:
            self.assertIn(f']({name})', text)
        self.assertIn('kommunale-haftpflicht/README.md', text)
        self.assertIn('Diese Testakte wurde mit KI generiert', text)
        rubric = yaml.safe_load((CASE/'rubric.yaml').read_text())
        self.assertEqual(rubric['plugin'], 'kommunale-haftpflicht')
        self.assertTrue(any('Absatz 8' in c['description'] for c in rubric['checks']))

    def test_combined_pdf(self):
        reader = PdfReader(PDF)
        self.assertGreaterEqual(len(reader.pages), 18)
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
        for needle in ('Ottokar Wolkenbein', 'Thusnelda Pimpinella', 'Cosima Konfetti', 'Gundula Funk', '403,20', '715,20'):
            self.assertIn(needle, text)
        self.assertNotIn('Diese Testakte wurde mit KI generiert', text)
        self.assertNotIn('This test case file was generated', text)

    def test_native_zip_bytes_and_flatness(self):
        path = PACKAGES/f'testakte-{SLUG}.zip'
        if not path.exists():
            self.skipTest('Archivprüfung benötigt VERLETZUNG_PAKETE oder dist')
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), len(set(names)))
            self.assertFalse(any('/' in n or '\\' in n for n in names))
            self.assertEqual(set(names), NAMES | {PDF.name, 'README.txt'})
            self.assertTrue(archive.read('README.txt').startswith(NOTICE_BYTES))
            for name in NAMES:
                self.assertEqual(archive.read(name), (CASE/name).read_bytes())
            self.assertEqual(archive.read(PDF.name), PDF.read_bytes())

    def test_individual_pdf_zip(self):
        path = PACKAGES/f'testakte-{SLUG}-einzelpdfs.zip'
        if not path.exists():
            self.skipTest('Archivprüfung benötigt VERLETZUNG_PAKETE oder dist')
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            self.assertEqual(set(names), {Path(n).stem+'.pdf' for n in NAMES} | {'README.txt'})
            self.assertFalse(any('/' in n or '\\' in n for n in names))
            self.assertTrue(archive.read('README.txt').startswith(NOTICE_BYTES))
            for name in names:
                if name.endswith('.pdf'):
                    reader = PdfReader(io.BytesIO(archive.read(name)))
                    text = '\n'.join(p.extract_text() or '' for p in reader.pages)
                    self.assertGreater(len(text), 200, name)
                    self.assertNotIn('Diese Testakte wurde mit KI generiert', text, name)
                    self.assertNotIn('This test case file was generated', text, name)


if __name__ == '__main__':
    unittest.main(verbosity=2)
