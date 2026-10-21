#!/usr/bin/env python3
"""Prüft Quellenpaket, eigenständige Prompts und bytegenaue Release-Zuordnung."""
from __future__ import annotations
import argparse
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import unittest
from urllib.parse import unquote
import zipfile
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from openpyxl import load_workbook
from pypdf import PdfReader
import yaml
from quality_lab import validate_profile
from testakte_disclaimer import NOTICE_BYTES

ROOT = Path(__file__).resolve().parents[1]
NAME = 'insolvenzsachbearbeitung'
PARENT = 'insolvenzforderungen-checker'
DIST = None


def module(filename):
    spec = importlib.util.spec_from_file_location(filename.replace('-', '_'), ROOT/'scripts'/filename)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class SachbearbeitungTest(unittest.TestCase):
    def test_four_cases_preserve_native_originals_and_real_email_attachments(self):
        cases = []
        for filename in ('insolvenzsachbearbeitung-betriebsakten.json', 'insolvenzsachbearbeitung-verbraucherakte.json'):
            cases += json.loads((ROOT/'scripts/data'/filename).read_text())
        self.assertEqual(len(cases), 4)
        for case in cases:
            directory = ROOT/'testakten'/case['slug']
            self.assertGreaterEqual(len(case['files']), 22)
            self.assertTrue({'.pdf', '.eml', '.docx', '.xlsx'} <= {Path(x['name']).suffix for x in case['files']})
            for item in case['files']:
                source = directory/item['name']
                self.assertTrue(source.is_file(), str(source))
                self.assertGreater(source.stat().st_size, 100)
                if source.suffix == '.eml':
                    message = BytesParser(policy=policy.default).parsebytes(source.read_bytes())
                    self.assertFalse(message.defects)
                    self.assertTrue(message['Date'])
                    actual = {part.get_filename(): part.get_payload(decode=True) for part in message.iter_attachments()}
                    expected = {name: (directory/name).read_bytes() for name in item.get('attachments', [])}
                    self.assertEqual(actual, expected, source.name)
                elif source.suffix == '.pdf':
                    text = '\n'.join(p.extract_text() or '' for p in PdfReader(source).pages)
                    self.assertGreater(len(text), 100)
                    self.assertNotIn('Diese Testakte wurde mit KI', text)
            if DIST is not None:
                for suffix in ('', '-einzelpdfs'):
                    with zipfile.ZipFile(DIST/f"testakte-{case['slug']}{suffix}.zip") as archive:
                        self.assertIsNone(archive.testzip())
                        self.assertEqual(len(archive.namelist()), len(case['files'])+1)
                        self.assertEqual(archive.namelist()[0], 'README.txt')
                        self.assertTrue(all('/' not in name for name in archive.namelist()))
                        self.assertTrue(archive.read('README.txt').startswith(NOTICE_BYTES))
                        if not suffix:
                            for item in case['files']:
                                self.assertEqual(archive.read(item['name']), (directory/item['name']).read_bytes())

    def test_consumer_workbooks_reconcile_without_double_counting_transfers(self):
        directory = ROOT/'testakten/insolvenzsachbearbeitung-familie-haberlein-nuernberg'
        checks = {
            '31_Haushalt.xlsx': [('Haushalt','B20',5018), ('Haushalt','C20',3097.4), ('Haushalt','B22',1920.6), ('Konten','E6',505.9), ('Konten','E7',-2120), ('Konten','E8',1591)],
            '32_Glaeubigerabgleich.xlsx': [('Abgleich','D11',277480), ('Abgleich','E11',275360), ('Abgleich','F11',271920)],
        }
        for filename, cells in checks.items():
            values = load_workbook(directory/filename, data_only=True)
            formulas = load_workbook(directory/filename, data_only=False)
            try:
                for name, cell, expected in cells:
                    self.assertAlmostEqual(values[name][cell].value, expected, places=2)
                    self.assertTrue(str(formulas[name][cell].value).startswith('='))
            finally:
                values.close(); formulas.close()
        # Unabhängige Belegrechnung: Überträge 4.438 EUR sind kein zusätzlicher Lohn.
        expenses = sum(map(Decimal, ('1440','760','420','140','80','79','95','83.40')))
        self.assertEqual(Decimal('5018')-expenses, Decimal('1920.60'))

    def test_consumer_dates_and_separate_proceedings_survive_pdf_export(self):
        directory = ROOT/'testakten/insolvenzsachbearbeitung-familie-haberlein-nuernberg'
        opening = '\n'.join(p.extract_text() or '' for p in PdfReader(directory/'27_Eroeffnung_Fridolin.pdf').pages)
        request = '\n'.join(p.extract_text() or '' for p in PdfReader(directory/'28_Verfuegung_Lale.pdf').pages)
        self.assertIn('9. September 2026', opening)
        self.assertIn('1. Oktober 2026', opening)
        self.assertRegex(opening, r'keine\s+Erteilung')
        self.assertIn('innerhalb eines Monats', request)
        self.assertIn('bisher nicht bestellt', request)
        self.assertIn('verkürzt', request)

    def test_eleven_skills_are_activatable_and_references_resolve(self):
        skills = sorted((ROOT/NAME/'skills').glob('*/SKILL.md'))
        self.assertEqual(len(skills), 11)
        for path in skills:
            text = path.read_text(encoding='utf-8')
            front = yaml.safe_load(text.split('---', 2)[1])
            self.assertEqual(set(front), {'name', 'description'})
            self.assertEqual(front['name'], path.parent.name)
            self.assertLessEqual(len(front['description']), 360)
            self.assertNotIn('§', front['description'])
            self.assertNotIn('im Plugin', front['description'])
            self.assertEqual(re.findall(r'^## ([1-6])\.', text, re.M), list('123456'))
            for word in ('Times New Roman', '11 pt', 'ausformuliert'):
                self.assertIn(word, text)
            for target in re.findall(r'\]\(([^)]+)\)', text):
                if target.startswith(('https:', 'http:', '#', 'mailto:')):
                    continue
                target = unquote(target.split('#')[0])
                self.assertTrue((path.parent/target).resolve().exists(), f'{path}: {target}')

    def test_prompts_are_separate_complete_and_bounded(self):
        builder = module('build-insolvenzsachbearbeitung-release.py')
        for source in builder.prompt_sources():
            self.assertEqual(source.read_bytes(), source.with_suffix('.txt').read_bytes())
            data = source.read_bytes().replace(b'\r\n', b'\n').replace(b'\n', b'\r\n')
            if 'werkstatt' not in source.name:
                self.assertLessEqual(len(data), 7500, source)
            else:
                self.assertGreater(len(source.read_text().split()), 6000 if source.parent.name == NAME else 2500)
            self.assertRegex(source.read_text(), r'(?i)trägt nicht|gilt nicht pauschal|grenze|daraus folgt keine|nicht jede')
        profile = json.loads((ROOT/'quality/evals'/f'{NAME}.json').read_text())
        validate_profile(profile, NAME, ROOT/NAME)
        self.assertGreaterEqual(len(profile['sources']), 8)
        self.assertGreaterEqual(len(profile['cases']), 4)

    def test_manifests_and_marketplace_agree(self):
        market = json.loads((ROOT/'.claude-plugin/marketplace.json').read_text())
        for name, version in ((NAME, '1.0.0'), (PARENT, '445.35.4')):
            entries = [p for p in market['plugins'] if p['name'] == name]
            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0]['version'], version)
            for relative in (('plugin.json', '.claude-plugin/plugin.json', '.codex-plugin/plugin.json') if name == NAME else ('.claude-plugin/plugin.json',)):
                manifest = json.loads((ROOT/name/relative).read_text())
                self.assertEqual((manifest['name'], manifest['version']), (name, version))
                self.assertEqual(manifest['description'], entries[0]['description'])

    def test_release_assets_and_packages_are_byte_exact(self):
        if DIST is None:
            self.skipTest('--dist erforderlich')
        builder = module('build-insolvenzsachbearbeitung-release.py')
        self.assertEqual({p.name for p in DIST.iterdir()}, builder.asset_names())
        lines = (DIST/'checksums-sha256.txt').read_text().splitlines()
        self.assertEqual(len(lines), len(builder.asset_names())-1)
        for line in lines:
            digest, name = line.split('  ', 1)
            self.assertEqual(sha(DIST/name), digest)
        for name in (NAME, PARENT):
            expected = builder.plugin_sources(name)
            for suffix, prefix in (('', ''), ('-portable', name+'/')):
                with zipfile.ZipFile(DIST/(name+suffix+'.zip')) as archive:
                    self.assertIsNone(archive.testzip())
                    self.assertEqual(set(archive.namelist()), {prefix+p for p in expected})
                    for relative, path in expected.items():
                        self.assertEqual(archive.read(prefix+relative), path.read_bytes())
                    for path in archive.namelist():
                        self.assertNotIn('..', Path(path).parts)
                        self.assertFalse(set(Path(path).parts) & {'testakte', 'testakten', 'prompts'})
        record = json.loads((DIST/'quellenabgleich.json').read_text())
        for source in builder.prompt_sources():
            asset = builder.prompt_asset(source)
            self.assertEqual((DIST/asset).read_bytes(), source.read_bytes())
            target = DIST/str(Path(asset).with_suffix('.pdf'))
            result = builder.inspect_pdf(source, target)
            for key, value in result.items():
                self.assertEqual(record['pdfs'][target.name][key], value)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dist', type=Path)
    args, remaining = parser.parse_known_args()
    DIST = args.dist.resolve() if args.dist else None
    unittest.main(argv=[__file__, *remaining])
