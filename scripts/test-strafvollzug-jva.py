#!/usr/bin/env python3
"""Integritätsprüfungen für Strafvollzug/JVA; Releaseprüfungen optional mit --dist."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import unittest
from urllib.parse import unquote
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader
import yaml

from release_routing import case_asset_url, plugin_asset_url
from testakte_disclaimer import NOTICE_BYTES, NOTICE_MARKDOWN, pdf_content_errors
from testakte_einzelpdf_common import document_arcname_pairs
from testakte_zip_common import working_dump_archive_pairs

ROOT = Path(__file__).resolve().parents[1]
NAME = "strafvollzug-jva"
PLUGIN = ROOT / NAME
DIST = None
LEGACY_CASES = (
    "jva-berlin-besuch-ernaehrung", "jva-nrw-unterbringung-gesundheit",
    "jva-bayern-bildung-computer", "jva-niedersachsen-lockerungen-familie",
)
EXTRA_CASES = ("jva-hessen-absonderung-hagedorn", "jva-bw-schutzmassnahmen-finkbeiner")
CASES = LEGACY_CASES + EXTRA_CASES
FILES = {
    "auftrag": "01_Auftrag_und_Verfahrensstand.docx",
    "antrag": "02_Antrag_an_die_JVA.docx",
    "bescheid": "03_Bescheid.pdf",
    "bekanntgabe": "04_Bekanntgabevermerk.pdf",
    "hausordnung": "05_Hausordnung_Auszug.pdf",
    "beleg1": "06_Sachbeleg.pdf",
    "beleg2": "07_Ergaenzende_Auskunft.pdf",
    "gespraech": "08_Gespraechsvermerk.docx",
    "gericht": "09_Gerichtsantrag_Entwurf.docx",
    "mail1": "10_Kontaktaufnahme.eml",
    "mail2": "11_Auskunft_der_JVA.eml",
    "mail3": "12_Externe_Auskunft.eml",
    "mail4": "13_Rueckfrage.eml",
    "excel": "14_Chronologie_Belege_Termine.xlsx",
    "mail5": "15_Nachricht_zum_Sachstand.eml",
    "plan": "16_Vollzugsplan_Auszug.pdf",
}
ANNEXES = dict(zip((f"A{i}" for i in range(1, 9)),
                  (FILES[k] for k in ("antrag", "bescheid", "bekanntgabe", "hausordnung",
                                      "beleg1", "beleg2", "gespraech", "plan"))))
KINDS = ("werkstatt", "schnellstart", "hauptproblem")
EXPECTED_DAYS = dict(zip(LEGACY_CASES, ([3, 8], [4], [14, 28], [10, 17, 35, 94])))


def module(filename):
    path = ROOT / "scripts" / filename
    spec = importlib.util.spec_from_file_location(path.stem.replace("-", "_"), path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def docx_text(path):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = []
    for paragraph in root.findall(".//w:p", ns):
        parts = []
        for node in paragraph.iter():
            if node.tag == "{" + ns["w"] + "}t":
                parts.append(node.text or "")
            elif node.tag in {"{" + ns["w"] + "}br", "{" + ns["w"] + "}tab"}:
                parts.append(" ")
        paragraphs.append("".join(parts))
    return "\n".join(paragraphs)


def pdf_text(data):
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)


def office_pdf_text(data):
    # LibreOffice marks ligatures with /ActualText. pypdf and pdfplumber can
    # duplicate the following character; Poppler respects the replacement text.
    executable = shutil.which("pdftotext")
    if not executable:
        raise RuntimeError("pdftotext (Poppler) ist für die native Office-PDF-Prüfung erforderlich.")
    return subprocess.run([executable, "-", "-"], input=data,
                          stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                          check=True).stdout.decode("utf-8")


def prompt_urls_and_prose(source, rendered):
    # The renderer wraps long URLs at arbitrary characters. Check the complete
    # URL independently, then keep the stricter word comparison for all prose.
    urls = re.findall(r"https?://[^\s)]+", source)
    for url in urls:
        pattern = r"\s*".join(re.escape(char) for char in url)
        rendered, count = re.subn(pattern, "", rendered, count=1)
        if count != 1:
            raise AssertionError(f"Unvollständige URL in der Prompt-Lesefassung: {url}")
        source = source.replace(url, "", 1)
    return source, rendered


def words(text):
    # Office may wrap immediately after an existing hyphen. Keep all letters,
    # digits and the hyphen; this only joins the layout whitespace.
    text = re.sub(r"(?<=\w)-\s+(?=\w)", "-", text)
    return Counter(re.findall(r"[\w§]+(?:-[\w§]+)*", text.casefold()))


def compact(text):
    return re.sub(r"\s+", "", text)


def printed_cell_value(cell):
    """Prüft den sichtbaren Zahlenwert, nicht binäre Gleitkomma-Artefakte."""
    value = cell.value
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and cell.number_format == "0.00":
        return f"{value:.2f}"
    return str(value)


def assets():
    result = {f"{NAME}.zip", f"{NAME}-portable.zip", "checksums-sha256.txt"}
    result.update(f"{NAME}-{kind}.{suffix}" for kind in KINDS for suffix in ("md", "txt", "pdf"))
    for case in CASES:
        result.update((f"testakte-{case}.zip", f"testakte-{case}-einzelpdfs.zip", f"{case}_gesamt.pdf"))
    return result


class JVAComponentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cases = json.loads((ROOT / "scripts/data/strafvollzug-jva-akten.json").read_text())

    def require_dist(self):
        if DIST is None:
            self.skipTest("Releaseprüfung erst mit --dist; Quellenprüfungen laufen vollständig.")
        self.assertTrue(DIST.is_dir(), str(DIST))

    def test_manifests_skills_commands_and_local_links(self):
        for rel in ("plugin.json", ".claude-plugin/plugin.json", ".codex-plugin/plugin.json"):
            obj = json.loads((PLUGIN / rel).read_text())
            self.assertEqual((obj["name"], obj["version"]), (NAME, "1.1.0"))
        catalog = json.loads((ROOT / ".claude-plugin/marketplace.json").read_text())
        entries = [p for p in catalog["plugins"] if p["name"] == NAME]
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0]["source"], "./" + NAME)
        skills = list((PLUGIN / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 11)
        for path in skills:
            with self.subTest(skill=path.parent.name):
                content = path.read_text()
                meta = yaml.safe_load(content.split("---", 2)[1])
                self.assertEqual(set(meta), {"name", "description"})
                self.assertEqual(meta["name"], path.parent.name)
                self.assertLessEqual(len(meta["description"]), 360)
                self.assertEqual(re.findall(r"(?m)^#{1,3} ([1-6])\. ", content), list("123456"))
                for required in ("Times New Roman", "11 pt", "ausformuliert"):
                    self.assertIn(required, content)
        commands = list((PLUGIN / "commands").glob("*.md"))
        self.assertEqual({p.stem for p in commands},
                         {"jva-fall", "jva-antrag", "jva-bescheid", "jva-eilrechtsschutz", "jva-pruefung"})
        for path in commands:
            parts = path.read_text().split("---", 2)
            self.assertTrue(yaml.safe_load(parts[1])["description"])
            self.assertIn("$ARGUMENTS", parts[2])
        for path in PLUGIN.rglob("*.md"):
            for target in re.findall(r"\]\(([^ )]+)\)", path.read_text()):
                if target.startswith(("https:", "http:", "#", "mailto:")):
                    continue
                dest = (path.parent / unquote(target.split("#")[0])).resolve()
                self.assertTrue(dest.is_relative_to(PLUGIN.resolve()), f"{path}: {target}")
                self.assertTrue(dest.exists(), f"{path}: {target}")

    def test_prompt_pairs_and_review_hashes(self):
        profile = json.loads((ROOT / "quality/evals/strafvollzug-jva.json").read_text())
        for kind, key in zip(KINDS, ("workshop_review", "mini_review", "focus_review")):
            path = PLUGIN / f"{NAME}-{kind}.md"
            self.assertEqual(path.read_bytes(), path.with_suffix(".txt").read_bytes())
            if kind != "werkstatt":
                self.assertLessEqual(path.stat().st_size, 7500)
            self.assertEqual(profile[key]["prompt_sha256" if kind == "hauptproblem" else "sha256"], sha(path))
        self.assertEqual(profile["focus_review"]["skill_sha256"],
                         sha(PLUGIN / "skills/strafvollzug-fall-steuern/SKILL.md"))

    def test_original_inventory_and_court_annexes(self):
        self.assertEqual([c["slug"] for c in self.cases], list(LEGACY_CASES))
        total = 0
        for c in self.cases:
            folder = ROOT / "testakten" / c["slug"]
            with self.subTest(case=c["slug"]):
                natives = {p.name for p in folder.iterdir() if p.suffix in {".pdf", ".docx", ".eml", ".xlsx"}}
                self.assertEqual(natives, set(FILES.values()))
                self.assertEqual(Counter(Path(n).suffix for n in natives),
                                 {".docx": 4, ".pdf": 6, ".eml": 5, ".xlsx": 1})
                self.assertEqual({p.name for p, _ in document_arcname_pairs(folder)}, natives)
                self.assertEqual({p.name for p, _ in working_dump_archive_pairs(folder, include_gesamt_pdf=False)}, natives)
                total += len(natives)
                for filename in natives:
                    path = folder / filename
                    self.assertFalse(path.is_symlink())
                    self.assertGreater(path.stat().st_size, 0)
                    if path.suffix in {".xlsx", ".docx"}:
                        with zipfile.ZipFile(path) as archive:
                            self.assertIsNone(archive.testzip())
                            self.assertIn("[Content_Types].xml", archive.namelist())
                draft = docx_text(folder / FILES["gericht"])
                self.assertIn(c["person"], draft)
                self.assertIn(c["record"], draft)
                self.assertIn("05.10.2026", draft)
                self.assertIn("16.10.2026", draft)
                self.assertIn("ein tatsächlich erfolgter gerichtlicher Eingang wird damit nicht behauptet", draft)
                self.assertEqual(set(re.findall(r"\bA\d+\b", draft)), set(ANNEXES))
                book = load_workbook(folder / FILES["excel"], data_only=True)
                rows = list(book["Belege"].iter_rows(min_row=7, max_row=14, values_only=True))
                self.assertEqual({r[0]: r[1] for r in rows}, ANNEXES)
                for label, filename in ANNEXES.items():
                    self.assertIn(label + ":", draft)
                    self.assertTrue((folder / filename).is_file())
                for path in folder.glob("*.pdf"):
                    data = path.read_bytes()
                    self.assertFalse(pdf_content_errors(data), path.name)
                    text = pdf_text(data)
                    self.assertGreater(len(text.split()), 40, path.name)
                    self.assertNotIn("\ufffd", text)
        self.assertEqual(total, 64)

    def test_email_headers_domains_ids_and_native_attachment_bytes(self):
        ids = set()
        attachment_count = 0
        for c in self.cases:
            folder = ROOT / "testakten" / c["slug"]
            for path in sorted(folder.glob("*.eml")):
                with self.subTest(mail=str(path.relative_to(ROOT))):
                    msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                    for part in msg.walk():
                        self.assertFalse(part.defects)
                    for header in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version"):
                        self.assertEqual(len(msg.get_all(header, [])), 1, header)
                    addresses = getaddresses([str(msg["From"]), str(msg["To"])])
                    self.assertTrue(addresses)
                    for _, address in addresses:
                        self.assertRegex(address, r"^[^@\s]+@[^@\s]+\.example$")
                    message_id = str(msg["Message-ID"])
                    self.assertNotIn(message_id, ids)
                    self.assertRegex(message_id, r"^<[^<>\s]+@[^<>\s]+\.example>$")
                    ids.add(message_id)
                    when = parsedate_to_datetime(str(msg["Date"]))
                    self.assertIsNotNone(when.utcoffset())
                    self.assertLessEqual(when.date(), date(2026, 10, 16))
                    body = msg.get_body(preferencelist=("plain",))
                    self.assertIsNotNone(body)
                    self.assertGreater(len(body.get_content().split()), 20)
                    attachments = list(msg.iter_attachments())
                    expected = {FILES["mail2"]: FILES["hausordnung"], FILES["mail3"]: FILES["beleg1"]}.get(path.name)
                    self.assertEqual([p.get_filename() for p in attachments], [expected] if expected else [])
                    for part in attachments:
                        self.assertEqual(part.get_content_type(), "application/pdf")
                        self.assertEqual(part.get_payload(decode=True), (folder / part.get_filename()).read_bytes())
                        attachment_count += 1
        self.assertEqual(len(ids), 20)
        self.assertEqual(attachment_count, 8)

    def test_workbook_dates_evidence_and_calculated_caches(self):
        for c in self.cases:
            path = ROOT / "testakten" / c["slug"] / FILES["excel"]
            before = sha(path)
            formulas = load_workbook(path, data_only=False)
            values = load_workbook(path, data_only=True)
            with self.subTest(case=c["slug"]):
                self.assertEqual(values.sheetnames, ["Chronologie", "Belege", "Termine"])
                for book in (formulas, values):
                    for sheet in book:
                        for row in sheet:
                            for cell in row:
                                self.assertNotEqual(cell.data_type, "e", f"{sheet.title}!{cell.coordinate}")
                for sheet in values:
                    self.assertEqual(sheet["B4"].value, datetime(2026, 10, 16))
                chrono = values["Chronologie"]
                populated = [r for r in chrono.iter_rows(min_row=7, values_only=True) if r[0] is not None]
                self.assertEqual(len(populated), len(c["dates"]))
                for row, expected in zip(populated, c["dates"]):
                    self.assertIsInstance(row[0], datetime)
                    self.assertEqual(row, (datetime.fromisoformat(expected[0]), expected[1], FILES[expected[2]], expected[3]))
                    self.assertTrue((path.parent / row[2]).is_file())
                future = [r for r in c["dates"] if r[0] >= "2026-10-16"]
                rows = [r for r in values["Termine"].iter_rows(min_row=7, values_only=True) if r[0] is not None]
                self.assertEqual(len(rows), len(future))
                self.assertEqual([r[3] for r in rows], EXPECTED_DAYS[c["slug"]])
                for index, (row, expected) in enumerate(zip(rows, future), 7):
                    self.assertEqual(row[0], datetime.fromisoformat(expected[0]))
                    self.assertEqual(row[1], expected[1] + "; " + FILES[expected[2]])
                    self.assertEqual(row[2], expected[3])
                    self.assertEqual(row[3], (row[0].date() - date(2026, 10, 16)).days)
                    self.assertEqual(formulas["Termine"][f"D{index}"].data_type, "f")
                    self.assertIsInstance(row[3], (int, float))
                self.assertEqual(sum(cell.data_type == "f" for sh in formulas for row in sh for cell in row), len(future))
                self.assertEqual(sha(path), before)

    def test_download_notices_links_and_component_routes(self):
        readme = (PLUGIN / "README.md").read_text()
        self.assertIn("/strafvollzug-jva-v1.1.0/", plugin_asset_url(NAME))
        for c in CASES:
            release = "/strafvollzug-jva-v1.0.0/" if c in LEGACY_CASES else "/strafvollzug-jva-v1.1.0/"
            text = (ROOT / "testakten" / c / "README.md").read_text()
            self.assertIn(NOTICE_MARKDOWN, text)
            for suffix in ("", "-einzelpdfs"):
                self.assertIn(release, case_asset_url(c, suffix))
                self.assertIn(f"testakte-{c}{suffix}.zip", text)
            self.assertIn(f"{c}_gesamt.pdf", text)
            self.assertLess(text.index(NOTICE_MARKDOWN), text.index("releases/download/"))
            self.assertIn(c + "/README.md", readme)
        for kind in KINDS:
            for suffix in ("md", "txt", "pdf"):
                self.assertIn(f"{NAME}-{kind}.{suffix}", readme)

    def test_extended_case_inventory_annexes_and_native_formats(self):
        report = json.loads((ROOT / "quality/strafvollzug-jva/erweiterungsakten-pruefung.json").read_text())
        self.assertEqual({c["slug"] for c in report["cases"]}, set(EXTRA_CASES))
        for case in report["cases"]:
            folder = ROOT / "testakten" / case["slug"]
            actual = {p.name: p for p, _ in working_dump_archive_pairs(folder, include_gesamt_pdf=False)}
            self.assertEqual(set(actual), {f["file"] for f in case["files"]})
            self.assertEqual(len(actual), 34)
            self.assertEqual(Counter(p.suffix for p in actual.values()),
                             {".docx": 8, ".pdf": 12, ".eml": 12, ".xlsx": 2})
            for item in case["files"]:
                path = actual[item["file"]]
                self.assertFalse(path.is_symlink())
                self.assertEqual(sha(path), item["sha256"], str(path))
                self.assertEqual(path.stat().st_size, item["bytes"])
                if path.suffix in {".docx", ".xlsx"}:
                    with zipfile.ZipFile(path) as archive:
                        self.assertIsNone(archive.testzip())
                        self.assertIn("[Content_Types].xml", archive.namelist())
                elif path.suffix == ".pdf":
                    self.assertFalse(pdf_content_errors(path.read_bytes()), path.name)
                    self.assertGreater(len(pdf_text(path.read_bytes()).split()), 60, path.name)
            draft = docx_text(actual[case["draft"]])
            self.assertGreater(len(draft.split()), 650)
            self.assertEqual(set(re.findall(r"\bA\d+\b", draft)), set(case["attachment_map"]))
            self.assertEqual(set(case["attachment_map"]), {f"A{i}" for i in range(1, 13)})
            book = load_workbook(actual["33_Chronologie_und_Anlagen.xlsx"], data_only=True)
            rows = list(book["Anlagen"].iter_rows(min_row=6, max_row=17, values_only=True))
            self.assertEqual({r[0]: r[1] for r in rows}, case["attachment_map"])
            for filename in case["attachment_map"].values():
                self.assertIn(filename, actual)

    def test_extended_emails_and_calculated_workbooks(self):
        report = json.loads((ROOT / "quality/strafvollzug-jva/erweiterungsakten-pruefung.json").read_text())
        ids = {str(BytesParser(policy=policy.default).parsebytes(p.read_bytes())["Message-ID"])
               for c in LEGACY_CASES for p in (ROOT / "testakten" / c).glob("*.eml")}
        attachments = Counter()
        for case in report["cases"]:
            folder = ROOT / "testakten" / case["slug"]
            self.assertEqual(len(case["mails"]), 12)
            for item in case["mails"]:
                msg = BytesParser(policy=policy.default).parsebytes((folder / item["file"]).read_bytes())
                for part in msg.walk():
                    self.assertFalse(part.defects)
                for header in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version"):
                    self.assertEqual(len(msg.get_all(header, [])), 1)
                for _, address in getaddresses([str(msg["From"]), str(msg["To"])]):
                    self.assertRegex(address, r"^[^@\s]+@[^@\s]+\.example$")
                self.assertNotIn(str(msg["Message-ID"]), ids)
                ids.add(str(msg["Message-ID"]))
                when = parsedate_to_datetime(str(msg["Date"]))
                self.assertIsNotNone(when.utcoffset())
                self.assertLessEqual(when.date(), date(2026, 10, 16))
                self.assertGreater(len(msg.get_body(preferencelist=("plain",)).get_content().split()), 45)
                attached = list(msg.iter_attachments())
                self.assertEqual([p.get_filename() for p in attached], [p["filename"] for p in item["attachments"]])
                for part, proof in zip(attached, item["attachments"]):
                    original = folder / proof["filename"]
                    self.assertEqual(part.get_payload(decode=True), original.read_bytes())
                    self.assertEqual(sha(original), proof["sha256"])
                    attachments[original.suffix] += 1
            self.assertEqual(len(case["workbooks"]), 2)
            for item in case["workbooks"]:
                path = folder / item["file"]
                formulas, values = load_workbook(path, data_only=False), load_workbook(path, data_only=True)
                self.assertEqual(values.sheetnames, [s["name"] for s in item["sheets"]])
                self.assertEqual(len(values.sheetnames), 4)
                checks = {(c["sheet"], c["cell"]): c["expected"] for c in item["formula_checks"]}
                self.assertTrue(checks)
                seen = set()
                for sheet in formulas:
                    for row in sheet:
                        for cell in row:
                            cached = values[sheet.title][cell.coordinate]
                            self.assertNotEqual(cached.data_type, "e")
                            if cell.data_type == "f":
                                key = (sheet.title, cell.coordinate)
                                self.assertIn(key, checks)
                                self.assertAlmostEqual(cached.value, checks[key], places=7)
                                seen.add(key)
                self.assertEqual(seen, set(checks))
                self.assertTrue(item["input_change_test"]["restored"])
                self.assertNotEqual(item["input_change_test"]["before"], item["input_change_test"]["after"])
        self.assertEqual(len(ids), 44)
        for suffix in (".pdf", ".docx", ".xlsx"):
            self.assertGreaterEqual(attachments[suffix], 2)

    def test_layout_word_joining_preserves_document_identifiers(self):
        self.assertEqual(words("LB- 2048/26"), words("LB-2048/26"))
        self.assertNotEqual(words("LB-2049/26"), words("LB-2048/26"))
        self.assertNotEqual(words("19.10.2026"), words("20.10.2026"))
        self.assertNotEqual(words("A1 A8"), words("A1"))
        source = "Beleg https://www.example.org/entscheidung_42.html bleibt maßgeblich."
        rendered = "Beleg https://www.exam\nple.org/entscheidung_\n42.html bleibt maßgeblich."
        left, right = prompt_urls_and_prose(source, rendered)
        self.assertEqual(words(left), words(right))
        with self.assertRaises(AssertionError):
            prompt_urls_and_prose(source, rendered.replace("42", "43"))
        left, right = prompt_urls_and_prose(source, rendered.replace("maßgeblich", ""))
        self.assertTrue(words(left) - words(right))

    def test_printed_decimal_values_respect_cell_precision(self):
        from openpyxl import Workbook
        cell = Workbook().active["A1"]
        cell.number_format = "0.00"
        cell.value = 10.26 - 1.55
        self.assertEqual(printed_cell_value(cell), "8.71")
        cell.value = 8.72
        self.assertNotEqual(printed_cell_value(cell), "8.71")
        cell.value = 14.2
        self.assertEqual(printed_cell_value(cell), "14.20")
        cell.number_format = "General"
        self.assertEqual(printed_cell_value(cell), "14.2")

    def test_dist_inventory_checksums_and_plugin_source_bytes(self):
        self.require_dist()
        expected = assets()
        self.assertEqual(len(expected), 30)
        self.assertEqual({p.name for p in DIST.iterdir()}, expected)
        checksums = {}
        for line in (DIST / "checksums-sha256.txt").read_text().splitlines():
            digest, name = line.split("  ", 1)
            self.assertRegex(digest, r"^[0-9a-f]{64}$")
            self.assertNotIn(name, checksums)
            checksums[name] = digest
        self.assertEqual(set(checksums), expected - {"checksums-sha256.txt"})
        for name, digest in checksums.items():
            self.assertEqual(digest, sha(DIST / name), name)
        sources = {p.relative_to(PLUGIN).as_posix(): p for p in PLUGIN.rglob("*")
                   if p.is_file() and p.name not in {"CLAUDE.md", ".DS_Store"}
                   and "__pycache__" not in p.parts and p.suffix != ".pyc"
                   and not (p.parent == PLUGIN and any(p.name == f"{NAME}-{k}.{s}" for k in KINDS for s in ("md", "txt", "pdf")))}
        sources["LICENSE"] = ROOT / "LICENSE"
        for suffix, prefix in (("", ""), ("-portable", NAME + "/")):
            with zipfile.ZipFile(DIST / f"{NAME}{suffix}.zip") as archive:
                self.assertIsNone(archive.testzip())
                self.assertEqual(set(archive.namelist()), {prefix + n for n in sources})
                self.assertEqual(len(archive.namelist()), len(sources))
                for name, path in sources.items():
                    self.assertEqual(archive.read(prefix + name), path.read_bytes(), name)
                for name in archive.namelist():
                    self.assertFalse(any(p in {"testakten", "__pycache__", "templates"} for p in Path(name).parts), name)
                    self.assertNotIn(Path(name).suffix, {".docx", ".xlsx", ".pdf", ".eml", ".pyc"})

    def test_dist_original_bytes_central_validators_and_exact_pdf_pages(self):
        self.require_dist()
        original_validator = module("validate-testakten-release-zips.py")
        single_validator = module("validate-testakten-einzelpdf-zips.py")
        total = 0
        for case in CASES:
            folder = ROOT / "testakten" / case
            original = DIST / f"testakte-{case}.zip"
            single = DIST / f"testakte-{case}-einzelpdfs.zip"
            original_validator.assert_same(original.name, original_validator.expected_entries(folder),
                                           original_validator.zip_entries(original, require_notice=True))
            single_validator.assert_same(single.name, ["README.txt", *single_validator.expected_arcnames(folder)],
                                         single_validator.zip_entries(single, expected_suffix=".pdf"))
            with zipfile.ZipFile(original) as archive:
                self.assertEqual(archive.namelist()[0], "README.txt")
                self.assertEqual(archive.read("README.txt"), NOTICE_BYTES)
                originals = list(working_dump_archive_pairs(folder, include_gesamt_pdf=False))
                self.assertEqual(len(archive.namelist()), len(originals) + 2)
                for original_path, filename in originals:
                    self.assertEqual(archive.read(filename), original_path.read_bytes(), f"{case}/{filename}")
                    total += 1
                combined_data = (DIST / f"{case}_gesamt.pdf").read_bytes()
                self.assertEqual(archive.read(f"{case}_gesamt.pdf"), combined_data)
                self.assertEqual((folder / "gesamt-pdf" / f"{case}_gesamt.pdf").read_bytes(), combined_data)
            combined = PdfReader(io.BytesIO(combined_data))
            page_number = 0
            with zipfile.ZipFile(single) as archive:
                self.assertEqual(archive.namelist()[0], "README.txt")
                self.assertEqual(archive.read("README.txt"), NOTICE_BYTES)
                self.assertEqual(len(archive.namelist()), len(document_arcname_pairs(folder)) + 1)
                for path, name in document_arcname_pairs(folder):
                    data = archive.read(name)
                    if path.suffix == ".pdf":
                        self.assertEqual(data, path.read_bytes(), name)
                    text = pdf_text(data)
                    self.assertNotIn("\ufffd", text, name)
                    self.assertNotIn("\u25a0", text, name)
                    if path.suffix == ".docx":
                        self.assertFalse(words(docx_text(path)) - words(text), f"{case}/{name}: Word-Text fehlt")
                    elif path.suffix == ".eml":
                        msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                        self.assertFalse(words(msg.get_body(preferencelist=("plain",)).get_content()) - words(text), name)
                    elif path.suffix == ".xlsx":
                        text = office_pdf_text(data)
                        for sheet in load_workbook(path, data_only=True):
                            for row in sheet:
                                for cell in row:
                                    if cell.value is None:
                                        continue
                                    value = printed_cell_value(cell)
                                    self.assertIn(compact(value), compact(text), f"{case}/{sheet.title}!{cell.coordinate}")
                    for page in PdfReader(io.BytesIO(data)).pages:
                        actual = combined.pages[page_number]
                        self.assertEqual(page.get_contents().get_data(), actual.get_contents().get_data(), (case, page_number))
                        self.assertEqual(tuple(page.mediabox), tuple(actual.mediabox))
                        self.assertEqual(page.extract_text(), actual.extract_text())
                        page_number += 1
            self.assertEqual(page_number, len(combined.pages))
        self.assertEqual(total, sum(len(working_dump_archive_pairs(ROOT / "testakten" / c, include_gesamt_pdf=False)) for c in CASES))

    def test_dist_prompt_pdf_full_text_and_source_identity(self):
        self.require_dist()
        for kind in KINDS:
            stem = f"{NAME}-{kind}"
            source = PLUGIN / f"{stem}.md"
            for suffix in ("md", "txt"):
                self.assertEqual((DIST / f"{stem}.{suffix}").read_bytes(), source.read_bytes())
            text = pdf_text((DIST / f"{stem}.pdf").read_bytes())
            prose, rendered_prose = prompt_urls_and_prose(source.read_text(), text)
            self.assertFalse(words(prose) - words(rendered_prose), kind)
            self.assertNotIn("\ufffd", text)
            self.assertNotIn("\u25a0", text)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dist", type=Path)
    args, remaining = parser.parse_known_args()
    DIST = args.dist
    unittest.main(argv=[sys.argv[0], *remaining])
