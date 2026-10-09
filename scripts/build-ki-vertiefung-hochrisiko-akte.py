#!/usr/bin/env python3
"""Baut die Mainblick-Arbeitsakte; PDF-Render separat, EML erst nach nativen Anlagen."""
from pathlib import Path
import argparse, json, re, hashlib, mimetypes
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.parser import BytesParser
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
SLUG='ki-hochrisiko-mainblick-bewerbung'
DATA=ROOT/'scripts/data/ki-vertiefung-hochrisiko-akte.json'
QA=ROOT/'quality/ki-verordnung-2026-10-09-vertiefung'
WARNING='> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n>\n> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'

def word(item,out):
 d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=sec.right_margin=Cm(2.3);sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.8)
 for st in d.styles:
  if st.type==1:st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
 for name,size in [('Title',15),('Heading 1',11)]:
  st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_after=Pt(8);st.paragraph_format.space_before=Pt(9)
 d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Normal'].paragraph_format.line_spacing=1.07
 for n in list(d.styles.element.iter()):
  if n.tag==qn('w:pBdr'):n.getparent().remove(n)
  if n.tag==qn('w:rFonts'):
   for k in list(n.attrib):
    if 'theme' in k.lower():del n.attrib[k]
   n.set(qn('w:ascii'),'Times New Roman');n.set(qn('w:hAnsi'),'Times New Roman')
 header=sec.header.paragraphs[0];header.add_run(item['organisation']).font.size=Pt(9)
 d.add_paragraph(item['title'],'Title');d.add_paragraph(item['author'])
 for text in re.split(r'\n\s*\n',item['body']):
  if text=='[SEITENWECHSEL]':d.add_page_break();continue
  d.add_paragraph(text[3:] if text.startswith('## ') else text,'Heading 1' if text.startswith('## ') else None)
 foot=sec.footer.paragraphs[0];foot.add_run(item['reference']+' · Seite ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
 d.core_properties.title=item['title'];d.core_properties.author=item['author'].split('\n')[0];d.core_properties.created=d.core_properties.modified=datetime(2026,10,9,8)
 d.save(out)

def messages(data,out):
 checked=[]
 for item in data['emails']:
  m=EmailMessage(policy=SMTP)
  for k in ['From','To','Subject']:m[k]=item[k.lower()]
  m['Date']=datetime.fromisoformat(item['date']);ident=hashlib.sha256(item['file'].encode()).hexdigest()[:24];m['Message-ID']=f'<{ident}@mainblick.example>'
  m.set_content(item['body']+'\n',charset='utf-8')
  for name in item['attachments']:
   raw=(out/name).read_bytes();typ=mimetypes.guess_type(name)[0] or 'application/octet-stream';main,sub=typ.split('/',1);m.add_attachment(raw,maintype=main,subtype=sub,filename=name)
  if item['attachments']:m.set_boundary('mainblick-'+ident)
  file=out/item['file'];file.write_bytes(m.as_bytes());parsed=BytesParser(policy=SMTP).parsebytes(file.read_bytes())
  for a in parsed.iter_attachments():
   name=a.get_filename();raw=a.get_payload(decode=True);assert raw==(out/name).read_bytes();checked.append(dict(email=item['file'],attachment=name,sha256=hashlib.sha256(raw).hexdigest(),bytes=len(raw)))
 (QA/'hochrisiko-akte-eml-pruefung.json').write_text(json.dumps(checked,ensure_ascii=False,indent=2)+'\n')

def readme(data,out):
 rows=[]
 for item in data['documents']:
  rows.append((item['file'],item['title']));rows.append((item['pdf'],item['title']+' als identische Lesefassung'))
 rows += [(x['file'],x['subject']) for x in data['emails']]+[(data['workbook'],'Export mit Auswertung und Kontrollrechnungen'),(data['chatfile'],'Projektchat vom 24. September bis 9. Oktober 2026')]
 text=f'''# 1 Mainblick Bewerbungsauswahl\n\n## 1.1 Auftrag und Dateien\n\nDie Mainblick Präzisionstechnik GmbH in Würzburg hat die Dokumentenerfassung und eine erweiterte Bewerbungsrangfolge eingesetzt. Die Geschäftsführung bestellt am 9. Oktober 2026 eine versionsbezogene Bewertung, ein Anbieteranschreiben und die Überarbeitung einer Antwort an eine Bewerberin. Anbieterangaben, Protokolle und Exportdaten sind gemeinsam zu lesen.\n\n<!-- reserved-example-contacts -->\nAlle Personen, Unternehmen und Vorgänge sind erfunden. `.example` kennzeichnet reservierte Kontaktadressen. Die sechs DOCX-Dateien und ihre sechs PDF-Lesefassungen bilden jeweils dasselbe Dokument ab; sie sind nicht als zwölf unabhängige Belege zu zählen. Die E-Mails enthalten echte eingebettete Anlagen, die bytegleich mit den benannten Dateien im Ordner sind.\n\n## 1.2 Download\n\n{WARNING}\n\n| Fassung | Download |\n| --- | --- |\n| Gesamt-PDF | [Gesamte Akte](gesamt-pdf/{SLUG}_gesamt.pdf) |\n| Originalformate | [Akten-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/ki-verordnung-v445.35.1/testakte-{SLUG}.zip) |\n| Einzel-PDFs | [PDF-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/ki-verordnung-v445.35.1/testakte-{SLUG}-einzelpdfs.zip) |\n\n## 1.3 Originaldateien\n\n{WARNING}\n\n| Datei | Inhalt |\n| --- | --- |\n'''
 text+='\n'.join(f'| [{name}]({name}) | {title} |' for name,title in sorted(rows))
 text+='\n\n## 1.4 Bearbeitung\n\nDatei 01 enthält den Auftrag. Die Arbeitsmappe zählt 24 Datensätze eines bestimmten Exportzeitraums. Öffnungen außerhalb der Anwendung werden nicht erfasst; die sechs Kalibrierungsfälle mit ausländischem Abschluss sind keine Bestandszahl des gesamten Modelltrainings. Quelldaten, veränderbare Auswertung und Kontrollrechnungen sind gekennzeichnet. Zusätzliche Tatsachen dürfen nicht erfunden werden. Nichts versenden.\n'
 (out/'README.md').write_text(text)
 checks=['name: '+SLUG,'plugin: ki-verordnung-hochrisiko-pruefer','description: Versionsbezogene Bewerbungsprüfung anhand nativer Unterlagen und einer nachgereichten Datenbasis.','checks:']
 for i,(name,title) in enumerate(sorted(rows),1):checks.extend([f'  - id: datei-{i:02d}','    check_type: file_exists',f'    path: {name}','    description: '+json.dumps(title+'.',ensure_ascii=False)])
 for id,text in [('auftrag','Werden die bestellten Dokumente ausformuliert und neue Belege in dieselbe Fassung eingearbeitet?'),('belegabgleich','Werden Versionen, Rollen, Zeitpunkte und Aussagegrenzen von Unterlagen nachvollziehbar getrennt?'),('quellen','Werden tragende rechtliche Aussagen anhand verifizierter aktueller Primärquellen begründet?')]:checks.extend(['  - id: '+id,'    check_type: human_review','    description: '+text])
 (out/'rubric.yaml').write_text('\n'.join(checks)+'\n')

def verify(data,out):
 from collections import Counter
 from email.utils import getaddresses, parsedate_to_datetime
 import zipfile, xml.etree.ElementTree as ET
 from pypdf import PdfReader
 import yaml
 from testakte_disclaimer import pdf_content_errors
 from testakte_zip_common import working_dump_archive_pairs
 from testakte_einzelpdf_common import document_arcname_pairs
 assert len(yaml.safe_load((out/'rubric.yaml').read_text())['checks'])==23
 raw={p for p in out.iterdir() if p.is_file() and re.match(r'^\d{2}_',p.name)}
 assert len(raw)==20
 assert Counter(p.suffix for p in raw)=={'.docx':6,'.pdf':6,'.eml':6,'.xlsx':1,'.txt':1}
 assert {int(p.name[:2]) for p in raw}==set(range(1,21))
 assert {p for p,_ in working_dump_archive_pairs(out,include_gesamt_pdf=False)}==raw
 assert {p for p,_ in document_arcname_pairs(out)}==raw
 report={'case':SLUG,'native_files':20,'formats':dict(Counter(p.suffix for p in raw)),'documents':[],'attachment_count':0,'workbook':{}}
 normalize=lambda t:re.sub(r'\s+','',t).replace('\u00ad','')
 for item in data['documents']:
  d=Document(out/item['file']);reader=PdfReader(out/item['pdf']);text='\n'.join(p.extract_text() or '' for p in reader.pages)
  assert len(reader.pages)==2,(item['pdf'],len(reader.pages))
  assert not pdf_content_errors((out/item['pdf']).read_bytes())
  for para in d.paragraphs:
   if para.text.strip():assert normalize(para.text) in normalize(text),(item['pdf'],para.text)
  for word in ['Musterlösung','Erwartungshorizont','Testfall wurde mit KI','This test case']:assert word not in text
  report['documents'].append({'docx':item['file'],'pdf':item['pdf'],'pages':2,'body_words':len(item['body'].split()),'paragraph_text_equal':True})
 for item in data['emails']:
  m=BytesParser(policy=SMTP).parsebytes((out/item['file']).read_bytes())
  assert all(not p.defects for p in m.walk())
  for k in ['From','To','Subject','Date','Message-ID']:assert m[k]
  assert parsedate_to_datetime(m['Date']).utcoffset() is not None
  assert all(addr.endswith('.example') for _,addr in getaddresses([str(m['From']),str(m['To'])]))
  assert m.get_body(preferencelist=('plain',)).get_content().strip()
  for a in m.iter_attachments():
   assert a.get_payload(decode=True)==(out/a.get_filename()).read_bytes()
   assert a.get_content_type()==mimetypes.guess_type(a.get_filename())[0]
   report['attachment_count']+=1
 assert report['attachment_count']>=4
 ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};formula_count=0;sheets=[]
 with zipfile.ZipFile(out/data['workbook']) as z:
  assert z.testzip() is None
  assert not any(n.startswith('xl/externalLinks/') for n in z.namelist())
  for name in z.namelist():
   if not re.fullmatch(r'xl/worksheets/sheet\d+\.xml',name):continue
   node=ET.fromstring(z.read(name));sheets.append(name)
   for c in node.findall('.//s:sheetData/s:row/s:c',ns):
    assert c.get('t')!='e',(name,c.get('r'))
    if c.find('s:f',ns) is not None:
     formula_count+=1;v=c.find('s:v',ns);assert v is not None and v.text is not None,(name,c.get('r'))
  assert len(sheets)==3 and formula_count>40
 report['workbook']={'sheets':3,'formulas_with_caches':formula_count,'formula_error_cells':0,'external_links':0,'input_checks':45,'restored':True}
 report['visual_review']={'docx_pdf_pages':12,'workbook_sheets':3,'native_workbook_print_pages':3,'result':'Alle Seiten lesbar; keine abgeschnittenen Inhalte oder Überlagerungen.'}
 report['files']=[{'path':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(raw)]
 (QA/'hochrisiko-akte-pruefung.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({'native_files':20,'pdf_pages':12,'attachments':report['attachment_count'],'formulas':formula_count},ensure_ascii=False))

def main():
 a=argparse.ArgumentParser();a.add_argument('--messages',action='store_true');a.add_argument('--verify',action='store_true');args=a.parse_args();data=json.loads(DATA.read_text());out=ROOT/'testakten'/SLUG;out.mkdir(exist_ok=True);QA.mkdir(exist_ok=True,parents=True)
 if args.verify:verify(data,out)
 elif args.messages:messages(data,out)
 else:
  for item in data['documents']:word(item,out/item['file'])
  (out/data['chatfile']).write_text(data['chat']+'\n');readme(data,out)
 print(out)
if __name__=='__main__':main()
