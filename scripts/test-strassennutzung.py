#!/usr/bin/env python3
"""Prüft Straßenrechts-Workflows und ihre beiden individuellen Arbeitsakten."""

from __future__ import annotations

import csv
from decimal import Decimal
from datetime import date, datetime, time
from email import policy
from email.parser import BytesParser
import hashlib
import json
from pathlib import Path
import re
import unittest

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageStat
from pypdf import PdfReader
import yaml

from readme_decimal_headings import normalize_decimal_headings
from testakte_zip_common import working_dump_flat_pairs


ROOT = Path(__file__).resolve().parent.parent
PLUGIN = ROOT / "strassennutzung-genehmigungen"
SCHULE = ROOT / "testakten/strassenverkehrsrecht-stvo-schulstrasse-lieferzone"
LINDENHOF = ROOT / "testakten/strassennutzung-poller-lieferzufahrt-lindenhof-muenster"
ORIGINALS = {
    "01_verkehrsanordnung_schulstrasse_buchenweg.docx": "ed855b9d5285beb8bf5e70300e99a1481eda457d79c716abe22ab0a033a7351b",
    "02_lieferfahrten_und_ausnahmegenehmigungen.xlsx": "7b63af6312c4b9e633eeb4a63912d6d1867d06883367b225bb9cde24e1397832",
    "03_email_baeckerei_ausnahmegenehmigung_2026-08-14.eml": "9cf932bccfdd0d10fcb7ea665389e4e17c6d7328c3969a2d7aedaa844ea62cb7",
    "04_bussgeldanhoerung_lieferwagen_2026-08-20.docx": "3240f4c945014aa3dcff1ee7e938e4d4c32efe5b227ec563e19cd83203bdb109",
    "05_verkehrszaehlung_buchenweg_roh.csv": "9cb9a96e28861465138f504363d88c00ce8d7856849ec1f7d23c3d48712197de",
    "06_email_elternvertretung_beobachtungen_2026-08-18.eml": "10b082ea55a666f491e12c7fb39d0f1ce34de3ab0600c2d4ec82f5c9bf77e447",
    "07_fahrernotiz_lieferung_2026-08-19.docx": "f1d26ea0bff655555da14aeaf9b219b57cda61a9e63690a8d15edb526203cbf7",
    "08_telefonvermerk_maessner_kruse_2026-08-21.docx": "d898e8d6ec675243c79615a13b8bd0bffcffac592e75b02cfb0588fdf60b025a",
    "09_email_disposition_intern_2026-08-19.eml": "ce2324f9b0cfd1db6d51a4701645fa6cecbf61a01e9e19e4b4c1970cec5f92ad",
    "10_zwischennachricht_stadt_ausnahmeantrag_2026-08-25.docx": "dee106d482daa58e6290d74079393fcc17a83b47d881b5f8e130582f99ae52f4",
}


def document_text(path):
    if path.suffix == ".docx":
        doc = Document(path)
        return "\n".join(
            [p.text for p in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )
    if path.suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if path.suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        body = message.get_body(preferencelist=("plain",))
        return body.get_content() if body else ""
    if path.suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return "\n".join(str(cell) for sheet in book for row in sheet.values for cell in row if cell is not None)
        finally:
            book.close()
    return path.read_text(encoding="utf-8-sig") if path.suffix in {".txt", ".csv"} else ""


class StrassennutzungTests(unittest.TestCase):
    def test_ten_discoverable_skills_and_manifest(self):
        manifest = json.loads((PLUGIN / ".claude-plugin/plugin.json").read_text())
        self.assertEqual(manifest["name"], PLUGIN.name)
        self.assertLessEqual(len(manifest["description"]), 300)
        skills = sorted((PLUGIN / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 10)
        for path in skills:
            with self.subTest(skill=path.parent.name):
                front = yaml.safe_load(path.read_text().split("---", 2)[1])
                self.assertEqual(set(front), {"name", "description"})
                self.assertEqual(front["name"], path.parent.name)
                self.assertRegex(front["name"], r"^[a-z0-9-]{1,64}$")
                self.assertGreaterEqual(len(front["description"]), 80)
                self.assertLessEqual(len(front["description"]), 1024)
                self.assertNotRegex(front["description"], r"\d,\d|[<>]")
                self.assertIn("zitierweise.md", path.read_text())

    def test_standalone_prompts_are_substantial_and_separate(self):
        for kind in ("schnellstart", "werkstatt"):
            path = PLUGIN / f"{PLUGIN.name}-{kind}.md"
            text = path.read_text()
            self.assertGreater(len(text.encode()), 3000)
            if kind == "schnellstart":
                self.assertLessEqual(len(text.encode()), 7500)
            else:
                self.assertGreater(len(text.encode()), 12000)
            self.assertNotIn(chr(167), text)
            self.assertRegex(text, r"BVerwG|OVG|Verwaltungsgericht")
            self.assertIn("StVO", text)
            self.assertRegex(text, r"Landes|Straßengesetz|Straßenrecht")
            self.assertTrue(all(re.match(r"\d+(?:\.\d+)*\.?\s", h)
                                for h in re.findall(r"^#{1,6}\s+(.+)$", text, re.M)))
        self.assertFalse(list((PLUGIN / "skills").rglob("*-werkstatt.md")))
        self.assertFalse(list((PLUGIN / "skills").rglob("*-schnellstart.md")))

    def test_plugin_relative_links_resolve(self):
        for path in PLUGIN.rglob("*.md"):
            text = path.read_text()
            for link in re.findall(r"\[[^\]]*\]\(([^)\s]+)\)", text):
                if link.startswith(("http:", "https:", "#", "mailto:")):
                    continue
                target = link.split("#", 1)[0]
                self.assertTrue((path.parent / target).exists(), (path, link))

    def test_uploaded_ten_sources_remain_byte_identical(self):
        for filename, digest in ORIGINALS.items():
            self.assertEqual(hashlib.sha256((SCHULE / filename).read_bytes()).hexdigest(), digest, filename)

    def test_both_cases_have_distinct_native_documents(self):
        for directory, minimum in ((SCHULE, 26), (LINDENHOF, 30)):
            pairs = working_dump_flat_pairs(directory, include_gesamt_pdf=False)
            self.assertGreaterEqual(len(pairs), minimum, directory.name)
            suffixes = {path.suffix for path, _ in pairs}
            self.assertTrue({".docx", ".eml", ".csv", ".pdf"} <= suffixes, suffixes)
            self.assertTrue(all("/" not in name and "\\" not in name for _, name in pairs))
            self.assertFalse(suffixes & {".md", ".yaml", ".yml"})
            self.assertEqual(len({name.casefold() for _, name in pairs}), len(pairs))

    def test_case_sources_have_no_solution_or_processing_markers(self):
        for directory in (SCHULE, LINDENHOF):
            for path, _ in working_dump_flat_pairs(directory, include_gesamt_pdf=False):
                text = document_text(path)
                for marker in ("Lösungsmatrix", "Musterlösung", "Testakte", "Formathinweis", "[TODO]", chr(167)):
                    self.assertNotIn(marker, text, path)

    def test_new_letters_are_substantial_a4_documents(self):
        for directory in (SCHULE, LINDENHOF):
            for path in directory.rglob("*.docx"):
                if directory == SCHULE and path.name in ORIGINALS:
                    continue
                text = document_text(path)
                self.assertGreaterEqual(len(text), 700, path)
                for section in Document(path).sections:
                    w, h = sorted((section.page_width.mm, section.page_height.mm))
                    self.assertAlmostEqual(w, 210, delta=1, msg=str(path))
                    self.assertAlmostEqual(h, 297, delta=1, msg=str(path))

    def test_mail_headers_and_readable_utf8_bodies(self):
        for directory in (SCHULE, LINDENHOF):
            for path in directory.rglob("*.eml"):
                message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                for header in ("From", "To", "Date", "Subject", "Message-ID", "Content-Type"):
                    self.assertTrue(message.get(header), (path, header))
                text = document_text(path)
                self.assertGreater(len(text), 350, path)
                self.assertNotIn("\ufffd", text, path)

    def test_csv_schemas_and_record_counts(self):
        for directory in (SCHULE, LINDENHOF):
            for path in directory.rglob("*.csv"):
                with path.open(encoding="utf-8-sig", newline="") as stream:
                    rows = list(csv.reader(stream, delimiter=";", strict=True))
                self.assertGreaterEqual(len(rows), 5, path)
                self.assertTrue(all(len(row) == len(rows[0]) for row in rows), path)
                self.assertEqual(len(set(rows[0])), len(rows[0]), path)

    def test_decimal_readme_navigation_is_idempotent(self):
        for directory in (PLUGIN, SCHULE, LINDENHOF):
            text = (directory / "README.md").read_text()
            self.assertIn("<!-- decimal-headings -->", text)
            self.assertEqual(normalize_decimal_headings(text), text, directory)

    def test_lindenhof_delivery_charges_reconcile_to_invoice(self):
        path = LINDENHOF / "23_fahrtenbuch_2026-08-24_bis_09-08.csv"
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream, delimiter=";"))
        self.assertEqual(len(rows), 24)
        self.assertEqual(len({row["Fahrt_ID"] for row in rows}), 24)
        amount = lambda row: Decimal(row["Zusatzentgelt_netto_EUR"].replace(",", "."))
        self.assertEqual(sum(map(amount, rows)), Decimal("192"))
        fresh = [row for row in rows if row["Betrieb"] == "Frischefahrt West" and amount(row)]
        drinks = [row for row in rows if row["Betrieb"] == "Getränke Röttger" and amount(row)]
        self.assertEqual((len(fresh), sum(map(amount, fresh))), (8, Decimal("144")))
        self.assertEqual((len(drinks), sum(map(amount, drinks))), (4, Decimal("48")))
        self.assertTrue(all("FW-260914-118" in row["Beleg"] for row in fresh))
        self.assertTrue(all("Ankündigung" in row["Beleg"] for row in drinks))
        invoice = document_text(LINDENHOF / "26_rechnung_frischefahrt_2026-09-14.pdf")
        for value in ("FW-260914-118", "144,00", "27,36", "171,36"):
            self.assertIn(value, invoice)

    def test_visual_attachments_are_nonblank_and_readable_size(self):
        for name in ("28_schluesselablage_lb17_2026-09-22.png", "29_chat_fruehdienst_2026-09-04.png"):
            with Image.open(LINDENHOF / name) as picture:
                self.assertGreaterEqual(picture.width, 900)
                self.assertGreaterEqual(picture.height, 1200)
                self.assertGreater(ImageStat.Stat(picture.convert("L")).stddev[0], 20)

    def test_workbook_preserves_csv_rows_and_numeric_formula_caches(self):
        path = LINDENHOF / "30_betriebsaufzeichnungen_lindenhof.xlsx"
        book = load_workbook(path, data_only=True)
        formulas = load_workbook(path, data_only=False)
        try:
            schemas = (
                ("Lieferfahrten", "23_fahrtenbuch_2026-08-24_bis_09-08.csv", 11),
                ("Schlüsselvorgänge", "24_schluesseljournal_2026-08-21_bis_09-22.csv", 9),
                ("Beobachtungen", "25_beobachtungen_lindenbogen_2026-08-24_bis_09-18.csv", 9),
            )
            self.assertEqual(book.sheetnames, [schema[0] for schema in schemas])
            for name, filename, columns in schemas:
                sheet = book[name]
                with (LINDENHOF / filename).open(encoding="utf-8-sig", newline="") as stream:
                    rows = list(csv.reader(stream, delimiter=";"))[1:]
                self.assertEqual(sheet.max_row, 7 + len(rows))
                self.assertEqual(str(sheet.page_setup.paperSize), "9")
                self.assertEqual(sheet.page_setup.orientation, "landscape")
                self.assertEqual(sheet.freeze_panes, "C8")
                for r, source in enumerate(rows, 8):
                    for c, expected in enumerate(source, 1):
                        self.assertLessEqual(c, columns)
                        value = sheet.cell(r, c).value
                        if isinstance(value, datetime):
                            value = value.strftime("%Y-%m-%d %H:%M") if " " in expected else value.date().isoformat()
                        elif isinstance(value, date):
                            value = value.isoformat()
                        elif isinstance(value, time):
                            value = value.strftime("%H:%M")
                        elif isinstance(value, (int, float)):
                            self.assertAlmostEqual(value, float(expected.replace(",", ".")), places=6)
                            continue
                        else:
                            value = "" if value is None else str(value)
                        self.assertEqual(value, expected, (name, r, c))
                for row in formulas[name]:
                    for cell in row:
                        if cell.data_type == "f":
                            cached = sheet[cell.coordinate].value
                            self.assertIsInstance(cached, (int, float), (name, cell.coordinate))
            self.assertEqual(book["Lieferfahrten"]["I4"].value, 192)
        finally:
            book.close()
            formulas.close()

    def test_source_pdfs_are_a4_and_not_empty(self):
        for directory in (SCHULE, LINDENHOF):
            for path in directory.glob("*.pdf"):
                pages = PdfReader(path).pages
                self.assertGreater(len(pages), 0, path)
                for page in pages:
                    dimensions = sorted((float(page.mediabox.width), float(page.mediabox.height)))
                    self.assertAlmostEqual(dimensions[0], 595.276, delta=2, msg=str(path))
                    self.assertAlmostEqual(dimensions[1], 841.89, delta=2, msg=str(path))
                    self.assertGreater(len(page.extract_text() or ""), 200, path)


if __name__ == "__main__":
    unittest.main()
