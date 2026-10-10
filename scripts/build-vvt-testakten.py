#!/usr/bin/env python3
"""Reproduzierbare fiktive Unterlagen für drei VVT-Arbeitsakten."""
from pathlib import Path
import argparse, copy, datetime as dt, json, uuid
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'scripts/data/vvt-testakten.json'
from testakte_disclaimer import NOTICE_MARKDOWN

WARNING=NOTICE_MARKDOWN
FIELDS=['purpose','joint_controller','controller_contacts','processing_categories','data_subjects','data_categories','recipients','transfers','retention','toms','legal_basis','systems','processor_contracts','sources','next_review']

def activity(id,title,role='controller',**kw):
    a=dict(id=id,title=title,activity_revision=1,role=role,owner='',status='aktiv',**{k:'' for k in FIELDS})
    a.update(kw)
    a['screening']=dict(high_risk='offen',art35_3='offen',positive_list='offen',severity='offen',likelihood='offen',criteria='',rationale='',sources='',priority='offen')
    a['review']=dict(decision='offen',reviewer='',reviewed_at='',rationale='',sources='',basis_sha256='')
    return a

def dataset():
    cases=[]
    cases.append(dict(
      slug='vvt-handwerk-finkenbeil-erfurt',label='Finkenbeil Holzbau in Erfurt',company='Finkenbeil Holzbau GmbH',address='Drosselhain 18, 99091 Erfurt',domain='finkenbeil.example',leader='Brunhilde Finkenbeil',contact='Walburga Klee',it='Milan Yılmaz',department='Rufus Morgenstern',recipient='Clara Winde, Datenschutzberatung, Mühlenbogen 7, 99084 Erfurt',email='clara.winde@beratung-winde.example',employees=24,
      profile=[
        'Die Finkenbeil Holzbau GmbH beschäftigt 24 Personen. Acht arbeiten in Geschäftsführung, Büro, Abrechnung und Einsatzplanung; 16 sind überwiegend auf Baustellen tätig. Zwölf Diensttelefone werden wechselnden Montageteams übergeben. Es ist noch nicht geklärt, welche Beschäftigten damit täglich personenbezogene Auftragsdaten bearbeiten.',
        'Lohnunterlagen werden monatlich an das Steuerbüro übermittelt. Das Büro führt laufend Kundentermine, Rechnungen und Gewährleistungsanfragen. Das alte Verzeichnis stammt aus April 2025. Frau Klee hat die neue Liste am 07.10.2026 vorbereitet, konnte aber Aufbewahrungsregeln und Dienstleisterangaben nicht vollständig ergänzen.',
        'Im November sollen sechs Montagefahrzeuge mit Ortung ausgerüstet werden. Parallel möchte die Personalverwaltung Bewerbungen über ein neues Portal entgegennehmen. Beide Vorhaben sind beauftragt zur Prüfung, aber noch nicht für den Produktiveinsatz freigegeben. Der Betrieb hat keine schriftlich dokumentierte Entscheidung zur Benennung einer Datenschutzbeauftragten vorgelegt.',
        'Bitte gleichen Sie das vorhandene Verzeichnis mit den Unterlagen ab, halten Sie offene Angaben fest und erstellen Sie einen bearbeitbaren Änderungsstand. Prüfen Sie die Rollen der Dienstleister und das weitere Vorgehen für Ortung und Bewerbungsportal getrennt. Stellen Sie Rückfragen so zusammen, dass Geschäftsführung, Büro und IT ihre jeweiligen Punkte beantworten können.'
      ],
      activities=[
        activity('VT-0001','Personalverwaltung und monatliche Lohnvorbereitung',owner='Walburga Klee',purpose='Beschäftigungsverhältnisse verwalten und Lohnabrechnung vorbereiten.',data_subjects='Beschäftigte und ausgeschiedene Beschäftigte.',data_categories='Kontaktdaten, Bankverbindung, Arbeitszeiten, Vergütung und Fehlzeiten.',recipients='Steuerbüro Anselm Lenz; weitere Empfängerkategorien noch nicht erhoben.',systems='Personalordner im Büro und Lohnvorbereitung im Büroprogramm.',toms='Abschließbarer Schrank; persönliche Bürokennungen. Rechteprüfung noch offen.',sources='01_Unternehmensprofil_und_Auftrag.docx; 06_TOM_Bestandsaufnahme.pdf'),
        activity('VT-0002','Auftragsabwicklung und Kundenbetreuung',owner='Rufus Morgenstern',purpose='Angebote, Baustellentermine, Aufträge und Gewährleistungsanfragen bearbeiten.',data_subjects='Private Auftraggeber, Ansprechpartner gewerblicher Kunden und Bewohner der Baustellenobjekte.',data_categories='Namen, Adressen, Telefonnummern, Rechnungsdaten und Fotos des Arbeitsorts.',recipients='Montageteams und beauftragte Nachunternehmer; Liste wird ergänzt.',systems='Büroprogramm, E-Mail und wechselnd zugeordnete Diensttelefone.',retention='Aufträge werden archiviert; eine dokumentierte Trennung der Dokumentarten fehlt.',transfers='unbekannt',sources='12_IT_2026-10-07.eml')
      ],
      change_activity=activity('VT-0003','Vorgesehenes Ortungsportal für Montagefahrzeuge',owner='Rufus Morgenstern',status='geplant',purpose='Geplant: Einsatzplanung und Wiederfinden von Fahrzeugen. Die Geschäftsführung möchte auch Standzeiten vergleichen; dazu gibt es noch keinen abgestimmten Zweck.',data_subjects='Fahrende Beschäftigte und bei Teamfahrten weitere Insassen.',data_categories='Fahrzeugkennung, Zeitstempel, Positionen und Zuordnung zum Team.',recipients='Dispositionsbüro; Zugang der Geschäftsführung ist vorgesehen.',transfers='unbekannt',retention='Anbieter nennt zwölf Monate als Voreinstellung; betriebliche Entscheidung fehlt.',systems='Spurspatz Cloud, sechs Fahrzeuge, Testkonto ohne angeschlossene Ortungsgeräte.',toms='Privatmodus wird angeboten, ist im Angebot aber nicht konfiguriert.',sources='08_Vorhaben_und_Aenderungen.docx; 10_Anbieterangebot.pdf'),
      changes='Die Ortung wird als gesonderte geplante Tätigkeit mit Status geplant erfasst. Der neue Eintrag ist keine Startfreigabe. Das Bewerbungsportal ist noch nicht beschrieben; es muss als eigener Vorgang erhoben und darf nicht still in die Lohnabrechnung eingeordnet werden.',
      avv=[
        'Walburga Klee hat am 06.10.2026 die vorhandenen Vertragsordner geprüft. Zum bisherigen Büroprogramm liegt eine unterschriebene Vereinbarung mit der Wolkengarn Cloudservice GmbH vor. Sie nennt Hosting und technischen Betrieb in Deutschland. Eine aktuelle Liste der weiteren eingesetzten Dienstleister war nicht beigefügt.',
        'Beim Steuerbüro liegt der Beratungsvertrag vor. Das Formular im alten Verzeichnis bezeichnet alle Empfänger pauschal als Auftragsverarbeiter. Frau Klee kann nicht sagen, wer diese Einordnung vorgenommen hat. Eine gesonderte Prüfung der tatsächlich erbrachten Leistungen wurde im Ordner nicht gefunden.',
        'Für Spurspatz und das Bewerbungsportal Lebenslaufgarten liegen Angebote vor. Eine unterzeichnete Vereinbarung, endgültige Leistungsbeschreibung oder bestätigte Unterauftragnehmerliste fehlt jeweils. Dieser Vermerk dokumentiert den gefundenen Bestand; er bestätigt keine rechtliche Freigabe.'
      ],
      tom=[
        'Milan Yılmaz hat am 07.10.2026 vier Bürogeräte und drei Diensttelefone stichprobenartig angesehen. Die Bürogeräte sind verschlüsselt. Zwei Kennungen waren noch ehemaligen Beschäftigten zugeordnet; ob darüber noch ein Zugriff möglich war, wurde nicht geprüft. Das zentrale E-Mail-Konto Disposition wird von drei Personen genutzt.',
        'Die Telefone sind durch eine PIN geschützt. Bei wechselnden Teams bleibt die bereits geöffnete Auftrags-App teilweise angemeldet. Eine Fernlöschung wurde eingerichtet, aber seit dem Gerätewechsel nicht erprobt. Baustellenfotos enthalten gelegentlich Hausbewohner im Hintergrund.',
        'Eine Sicherung des Büroprogramms läuft täglich. Der letzte dokumentierte Wiederherstellungstest stammt aus Februar 2025. Für das neue Ortungsportal sind Rollen, Protokollierung und Privatmodus noch nicht eingestellt. Keine Ortungseinheit wurde bisher verbunden.'
      ],
      retention=[
        'Die Ordner werden nach Auftragsjahr archiviert. Angebote ohne Auftrag, Fotos, Rechnungen und Gewährleistungsunterlagen liegen im selben Verzeichnis. Eine Löschroutine nach Dokumentart ist nicht eingerichtet. Alte Bewerbungen befinden sich zum Teil im allgemeinen E-Mail-Postfach.',
        'Frau Klee möchte künftig Zuständigkeit, Beginn einer Aufbewahrungsphase und tatsächlich ausgeführte Löschungen festhalten. Dafür fehlen bislang Entscheidungen der Fachbereiche und ein Abgleich mit den aufzubewahrenden Dokumenten. Der vorliegende Arbeitsstand legt keine allgemeine Aufbewahrungsdauer für sämtliche Unterlagen fest.',
        'Der Anbieter Spurspatz nennt zwölf Monate als technische Voreinstellung. Diese Angabe wurde bisher nicht als betriebliche Löschentscheidung übernommen. Für Standortdaten und das Bewerbungsportal ist eine eigene Entscheidung vorgesehen.'
      ],
      proposal=[
        'Spurspatz Mobilitätssoftware GmbH, Antennenhain 5, 04129 Leipzig | Angebot SP-261002 vom 02.10.2026 an Finkenbeil Holzbau GmbH',
        'Wir bieten sechs Ortungseinheiten und das zugehörige Portal zur monatlichen Pauschale von 114,00 EUR netto an. Das Portal speichert Zeitstempel und Fahrzeugpositionen während der angeschalteten Zündung. In der Standardkonfiguration werden die Datensätze zwölf Monate vorgehalten.',
        'Ein Privatmodus und eingeschränkte Dispositionsrollen können zusätzlich konfiguriert werden. Eine persönliche Auswertung von Standzeiten ist technisch möglich. Diese Optionen sind im Testkonto noch nicht festgelegt. Wir benötigen Ihre Entscheidung zu den jeweiligen Einstellungen vor der Anbindung der Fahrzeuge.',
        'Das Rechenzentrum liegt nach unserer derzeitigen Produktbeschreibung in Deutschland. Die Liste der Supportunternehmen erhalten Sie mit den Vertragsunterlagen. Eine verbindliche Aussage zu einzelnen Supportzugriffen ist diesem Angebot nicht beigefügt. Ein Produktivstart ist für den 02.11.2026 angefragt, aber noch nicht bestätigt.'
      ],
      initiative=[
        'Besprechung am 06.10.2026 zwischen Brunhilde Finkenbeil, Rufus Morgenstern und Walburga Klee. Die Fahrzeuggeräte liegen noch ungeöffnet im Büro. Der Testzugang enthält bisher nur erfundene Fahrzeugkennungen.',
        'Rufus möchte den nächsten freien Transporter erkennen und Kundinnen bei Verspätung informieren. Brunhilde möchte zusätzlich wissen, warum manche Teams länger am Großhandel stehen. Es wurde noch nicht festgelegt, ob eine Auswertung je Person erfolgen soll und wie Privatfahrten behandelt werden.',
        'Walburga möchte das Bewerbungsportal Lebenslaufgarten einführen. Der Anbieter bewirbt eine Sortierung nach passenden Qualifikationen. Ob hierfür eine automatische Bewertung eingesetzt wird oder nur von Bewerbenden gesetzte Filter genutzt werden, hat sie noch nicht erfragt.',
        'Vor dem nächsten Termin sollen die jeweiligen Zwecke, Einstellungen, Nutzergruppen und benötigten Nachweise gesammelt werden. Weder eine Freigabe zur Beschäftigtenauswertung noch eine Freigabe zur Verarbeitung echter Bewerbungen wurde in der Besprechung erteilt.'
      ],
      requests=[
        'bitte ergänzen Sie bis zum 19.10.2026 die tatsächlichen Verarbeitungsschritte unserer Personalverwaltung und Auftragsabwicklung. Benötigt werden die Empfänger, die eingesetzten Programme, die zuständigen Personen und die vorhandenen Entscheidungen zur Aufbewahrung. Bitte teilen Sie außerdem mit, welche Beschäftigten regelmäßig selbst mit personenbezogenen Daten arbeiten; die bloße Zahl der Geräte genügt dafür nicht.',
        'Für Spurspatz benötige ich vor dem Anschluss der Fahrzeuge die Zwecke der Auswertungen, Angaben zu Privatfahrten, den geplanten Personenzuordnungen und zu sämtlichen Zugriffsrechten. Bitte legen Sie die endgültige Vertragsfassung und die Angaben zu Supportunternehmen und Supportorten bei. Die voreingestellte Speicherdauer wurde bislang nicht betrieblich beschlossen.',
        'Bitte beschreiben Sie das Bewerbungsportal gesondert. Insbesondere muss verständlich werden, ob eine Sortierung nur nach gewählten Kriterien erfolgt oder ob Bewerbende automatisch bewertet werden. Die Unterlagen sollen eine konkrete Prüfung ermöglichen; dieses Schreiben gibt keinen der beiden geplanten Dienste frei.'
      ],
      mails=[
        ('11_GF_2026-10-06.eml','2026-10-06','Brunhilde Finkenbeil','geschaeftsfuehrung','Wir sind doch nur 24 Personen','Guten Tag Frau Winde, im alten Ordner liegt eine Notiz, wonach unter 250 Leuten kein Verzeichnis gebraucht werde. Gleichzeitig rechnen wir jeden Monat Löhne ab und bearbeiten täglich Kundenaufträge. Bitte schauen Sie sich das konkret an. Ich möchte eine Liste, mit der Frau Klee weiterarbeiten kann, und keine pauschale Aussage nur nach unserer Größe.\n\nDie neuen Ortungsgeräte bleiben bis zur Klärung im Karton. Herr Morgenstern soll seine Zwecke genauer beschreiben. Wir haben noch keinen Termin für eine Beschäftigteninformation festgesetzt. Freundliche Grüße, Brunhilde Finkenbeil',None),
        ('12_IT_2026-10-07.eml','2026-10-07','Milan Yılmaz','it','Gerätebestand und offene Zugänge','Guten Tag Frau Winde, mein erster Rundgang ist dokumentiert. Zwölf Mobiltelefone werden von mehreren Teams geteilt; eine Liste der tatsächlich nutzenden Personen fehlt. Drei Leute greifen auf das gemeinsame Dispositionspostfach zu. Bei den alten Bürokennungen muss ich noch prüfen, ob sie nur angezeigt oder wirklich nutzbar sind.\n\nDie Behauptung alle Daten nur im Büro trifft jedenfalls für Fotos und E-Mails nicht zu. Das Büroprogramm wird gehostet. Die neue Ortung läuft noch nicht. Ich hänge meine Bestandsaufnahme an und ergänze die Supportorte, sobald die Anbieter geantwortet haben. Grüße, Milan Yılmaz','06_TOM_Bestandsaufnahme.pdf'),
        ('13_Fachbereich_2026-10-08.eml','2026-10-08','Rufus Morgenstern','disposition','Wofür ich die Positionen brauche','Guten Tag Frau Winde, mir würde meist die aktuelle Position des Fahrzeugs genügen. Ich muss sehen, wer einen Notauftrag in der Nähe übernehmen kann. Die Chefin sprach dagegen von Monatsvergleichen der Standzeiten. Ob dabei Teamnamen oder einzelne Fahrer erscheinen sollen, haben wir noch nicht entschieden.\n\nZwei Transporter dürfen nach Absprache mit nach Hause genommen werden. Ich kann nicht garantieren, dass außerhalb einer Baustellenfahrt keine privaten Wege vorkommen. Bitte behandeln Sie meine technische Wunschliste noch nicht als fertige Vorgabe für den Betrieb. Mit freundlichen Grüßen, Rufus Morgenstern',None),
        ('14_Anbieter_2026-10-09.eml','2026-10-09','Hedwig Funke','anbieter','Konfiguration Spurspatz','Guten Tag Frau Winde, Frau Klee hat uns Ihre Anfrage weitergeleitet. Die Funktion Privatmodus wird erst nach Ihrer Auswahl eingerichtet. Im Testkonto ist kein Fahrzeug angebunden. Das Angebot im Anhang nennt deshalb nur die verfügbare Standardkonfiguration.\n\nUnsere Vertragsabteilung stellt die aktuelle Dienstleisterliste zusammen. Die Formulierung Rechenzentrum Deutschland beantwortet nicht jede Frage zum Supportzugriff. Ich reiche diese Angaben nach. Bitte nennen Sie uns den vorgesehenen Nutzerkreis und ob eine Zuordnung zu einzelnen Beschäftigten gewünscht ist. Freundliche Grüße, Hedwig Funke','10_Anbieterangebot.pdf')
      ]))
    cases.append(dict(
      slug='vvt-cloudservice-wolkengarn-berlin',label='Wolkengarn Cloudservice in Berlin',company='Wolkengarn Cloudservice GmbH',address='Wolkensteg 22, 10997 Berlin',domain='wolkengarn.example',leader='Bruno Wolk',contact='Juno Nguyễn',it='Sami Arslan',department='Walburga Riedel',recipient='Mira Kranich, Datenschutzberatung, Fadenhof 3, 10247 Berlin',email='mira.kranich@kranich-beratung.example',employees=12,
      profile=[
        'Die Wolkengarn Cloudservice GmbH beschäftigt zwölf Personen und betreibt Dokumentenräume für kleine Geschäftskunden. Sie hostet Dateien, verwaltet Zugänge nach Kundenanweisung und unterstützt bei technischen Problemen. Über die Inhalte der hochgeladenen Dokumente entscheidet nach den bisherigen Verträgen der jeweilige Kunde.',
        'Daneben führt Wolkengarn eigene Personalakten, Lohnvorbereitung und Vertriebsdaten. Das bisherige Tabellenblatt mischt diese eigenen Aufgaben mit den Tätigkeiten für Kunden. Als Kundenkontakte sind bisher Finkenbeil Holzbau GmbH in Erfurt und Lampenhof Versand GmbH in Potsdam genannt; die vollständige Kundenliste liegt im Vertrieb.',
        'Ein neuer Supportdienst soll bei nächtlichen Störungen helfen. Die Geschäftsführung spricht von einer EU-Cloud, während die technische Anbieterantwort Supportteams in Toronto und Bengaluru erwähnt. Der geplante Zugriff ist noch nicht eingerichtet. Unklar sind Zugriffsumfang, technische Abschirmung und vertragliche Zuordnung.',
        'Bitte erstellen Sie aus dem gemischten Bestand nachvollziehbar getrennte Verzeichniseinträge für die eigenen Tätigkeiten und die Tätigkeiten für Auftraggeber. Dokumentieren Sie belegte Änderungen und fehlende Angaben. Die Entscheidung über den neuen Supportdienst ist ausdrücklich noch offen.'
      ],
      activities=[
        activity('VT-0001','Hosting und technischer Betrieb von Kundenräumen','processor',owner='Sami Arslan',controller_contacts='Finkenbeil Holzbau GmbH, Drosselhain 18, 99091 Erfurt, Walburga Klee, buero@finkenbeil.example. Lampenhof Versand GmbH, Pappelhof 8, 14482 Potsdam, Emma Lux, datenschutz@lampenhof.example. Weitere Kundenkontakte fehlen.',processing_categories='Speicherung, Bereitstellung, Sicherung und weisungsgebundene technische Unterstützung.',data_categories='Dokumente und Benutzerkonten nach Kundenauftrag; einzelne Auftragskategorien noch nicht erhoben.',transfers='Rechenzentren Deutschland und Finnland; Supportorte nicht abschließend erhoben.',toms='Mandantentrennung; persönliche Administrationskonten; Zugang wird protokolliert.',systems='Wolkengarn Kundenraum, Sicherungsdienst Garnspeicher.',sources='05_AVV_Vertragsbestand.pdf; 06_TOM_Bestandsaufnahme.pdf'),
        activity('VT-0002','Eigene Personalverwaltung und Lohnvorbereitung',owner='Juno Nguyễn',purpose='Eigene Beschäftigungsverhältnisse verwalten und monatliche Abrechnung vorbereiten.',data_subjects='Beschäftigte und Bewerbende; Abgrenzung der Bewerbungsphase noch offen.',data_categories='Stammdaten, Kontaktdaten, Vergütung, Bankverbindung und Arbeitszeiten.',recipients='Steuerkanzlei Ottokar Sand; interner Zugriff durch Personalverwaltung.',systems='Wolkengarn Personalordner und Lohnportal.',sources='01_Unternehmensprofil_und_Auftrag.docx; 13_Fachbereich_2026-10-08.eml')
      ],
      change_activity=None,
      changes='Beim Auftragsverarbeitereintrag werden die neuen Anbieterangaben zu möglichen Supportzugriffen ergänzt. Der Zusatz bezeichnet ein geprüft werden sollendes Vorhaben; er bestätigt weder eine Beauftragung noch eine zulässige Übermittlung. Die eigene Personalverwaltung bleibt ein eigener Verantwortlicheneintrag.',
      avv=[
        'Juno Nguyễn hat am 06.10.2026 die Verträge der beiden zuerst angelegten Kundenräume herausgesucht. Beide Kunden beauftragen das Speichern, Bereitstellen und Sichern ihrer Dokumente. Die Kunden legen Nutzerberechtigungen und Löschaufträge fest. Wolkengarn darf die Inhalte nach dem vorliegenden Wortlaut nicht für eigene Produktanalysen verwenden.',
        'Die bisherige Anlage zu Unterauftragnehmern nennt Garnspeicher Rechenzentrum GmbH in Deutschland und Suomen Lankadata Oy in Finnland. Die Anlage trägt den Stand 01.02.2025. Ein neuer Supportanbieter ist darin nicht genannt. Die Vertragsänderung für den Support liegt nur als nicht unterzeichneter Entwurf vor.',
        'Die weitere Kundschaft ist im CRM erfasst. Die vorgelegte Liste enthält nur zwei Auftraggeberkontakte und ist erkennbar unvollständig. Bei einem Kunden wurde ein Datenschutzkontakt aus einer alten Signatur übernommen. Seine Zuständigkeit ist noch zu bestätigen.'
      ],
      tom=[
        'Sami Arslan hat am 07.10.2026 die Administrationswege beschrieben. Produktivzugänge verlangen einen zweiten Faktor. Administrationsrechte werden personenbezogen vergeben. Freigeschaltete Rechte werden bei Ausscheiden entzogen; eine regelmäßige Übersicht über Änderungen wird jedoch nicht dauerhaft archiviert.',
        'Der technische Support kann nach einer gesonderten Kundenfreischaltung zeitweise Dokumenteninhalte sehen. Eine vollständig auf Metadaten beschränkte Unterstützung ist bei Fehlern in hochgeladenen Dateien nicht immer möglich. Protokolle werden in einem eigenen System gespeichert; die tatsächlich konfigurierte Dauer wird noch geprüft.',
        'Für den neuen Nachtsupport ist weder ein Administrationskonto noch ein Fernzugangsprofil angelegt. Das Anbieterangebot nennt eine Möglichkeit zeitlich begrenzter Sitzungen. Welche Personen aus welchen Ländern zugreifen und wer eine Sitzung freigibt, ist noch nicht dokumentiert.'
      ],
      retention=[
        'Kunden können Dokumente im Portal löschen. Die technische Beschreibung nennt eine nachgelagerte Sicherung, aus der einzelne Dokumente nicht sofort separat entfernt werden. Die Dauer bis zur Überschreibung ist im Betriebsprotokoll noch zu bestätigen. Eine beim Kunden sichtbare Löschung ist deshalb nicht ohne Weiteres die vollständige Entfernung sämtlicher Kopien.',
        'Nach Vertragsende exportiert Wolkengarn auf Anforderung den Kundenraum und sperrt die Benutzerkonten. Die bisherige Checkliste enthält kein ausgefülltes Feld zum Abschluss der Löschung im Sicherungsdienst. Für den Fall streitiger Zahlungsansprüche existiert noch keine nach Datenarten getrennte Entscheidung.',
        'Personalunterlagen von Wolkengarn werden außerhalb der Kundenräume geführt. Ihre Aufbewahrung folgt nicht dem Löschauftrag eines Hostingkunden. Die verantwortliche Person und die dokumentierten Regeln sind hierfür gesondert zu erheben.'
      ],
      proposal=[
        'Nachtfaden Support Ltd., Servicekontakt support@nachtfaden.example | Angebot NF-261003 vom 03.10.2026 an Wolkengarn Cloudservice GmbH',
        'Wir bieten technischen Bereitschaftsdienst für kritische Störungen an. Die Supportkoordination erfolgt aus Dublin. Je nach Verfügbarkeit können Spezialisten aus Toronto oder Bengaluru zugeschaltet werden. Die Server Ihrer Kundenräume sollen an ihren bisherigen Standorten verbleiben.',
        'Für jede Sitzung kann eine zeitliche Begrenzung eingerichtet werden. Ob die Spezialisten nur Systemprotokolle oder auch Dokumenteninhalte sehen müssen, hängt vom Fehlerbild ab. Die genaue Berechtigungsmatrix und das Freigabeverfahren sind mit Ihrem Betrieb abzustimmen.',
        'Ein Entwurf der vertraglichen Anlagen wird nach Ihrer Rückmeldung bereitgestellt. Dieses Angebot legt noch keine abschließende rechtliche Zuordnung der Beteiligten und keine dokumentierte Grundlage für internationale Zugriffe fest. Ein technischer Test mit erfundenen Daten ist möglich.'
      ],
      initiative=[
        'Bruno Wolk und Sami Arslan haben am 06.10.2026 den Bereitschaftsdienst besprochen. Es soll zunächst ausschließlich ein Test mit erfundenen Dateien stattfinden. Kundeninhalte dürfen noch nicht in das Testsystem übertragen werden.',
        'Bruno ging aufgrund der bisherigen Vertriebsfolie davon aus, dass sich sämtliche Verarbeitung in der EU abspielt. Sami wies auf die Anbieterangaben zu Toronto und Bengaluru hin. Die Folie wurde daraufhin für neue Angebote zurückgestellt; eine neue Aussage wurde noch nicht abgestimmt.',
        'Vor der weiteren Entscheidung sind die Auftraggeberrollen, Supportorte, Zugriffsmöglichkeiten und Vertragsanlagen zu klären. Die Einordnung des neuen Dienstes darf nicht allein aus dem Speicherort der Server abgeleitet werden. Das Projekt bleibt bis zu dieser Prüfung ohne Produktivzugang.'
      ],
      requests=[
        'bitte stellen Sie mir bis zum 19.10.2026 die vollständige Liste unserer Auftraggeber mit den jeweiligen Kontakten und den für sie tatsächlich ausgeführten Verarbeitungskategorien bereit. Die zwei bislang genannten Kunden sind nach Auskunft des Vertriebs nur ein Ausschnitt. Bitte trennen Sie diese Angaben von unserer eigenen Personalverwaltung und unseren eigenen Vertriebszwecken.',
        'Zum geplanten Nachtsupport benötige ich die eingesetzten Gesellschaften, die tatsächlichen Arbeitsorte der Supportpersonen, die Berechtigungsmatrix und die vorgesehenen Zugriffssperren. Bitte legen Sie Vertragsentwurf, Dienstleisterliste und die verfügbaren Unterlagen zu internationalen Zugriffen bei. Die Angabe EU-Hosting genügt nicht als Beschreibung eines möglichen Fernzugriffs.',
        'Bitte erläutern Sie außerdem, wann eine im Kundenportal gelöschte Datei auch aus Sicherungen verschwindet und wie dieser Vorgang bei Vertragsende nachgewiesen wird. Der nächste Registerstand soll die belegten Informationen und die verbleibenden Lücken erkennen lassen. Eine Freigabe des neuen Dienstleisters ist mit dieser Anfrage nicht verbunden.'
      ],
      mails=[
        ('11_GF_2026-10-06.eml','2026-10-06','Bruno Wolk','geschaeftsfuehrung','Welche Rolle haben wir eigentlich','Guten Tag Frau Kranich, wir schreiben in Angeboten, dass wir Auftragsverarbeiter sind. Beim Blick in die Tabelle finde ich aber auch unsere eigenen Personalakten und Vertriebskontakte unter derselben Überschrift. Bitte helfen Sie uns, das auseinanderzuhalten. Wir möchten einen Stand, den wir anschließend selbst pflegen können.\n\nBeim Nachtsupport hatte ich EU-Hosting mit ausschließlich europäischen Zugriffen gleichgesetzt. Sami hat mir die andere Anbieterantwort gezeigt. Bitte gehen Sie von einem noch nicht entschiedenen Vorhaben aus. Freundliche Grüße, Bruno Wolk',None),
        ('12_IT_2026-10-07.eml','2026-10-07','Sami Arslan','it','Support sieht nicht immer nur Metadaten','Guten Tag Frau Kranich, bei manchen Fehlern müssen wir die betroffene Datei öffnen. Dafür gibt es eine Kundenfreischaltung. Das ist etwas anderes als die reine Serverüberwachung. Beim neuen Anbieter sollen Zugriffe erst nach Freigabe möglich sein, aber die technische Konfiguration fehlt.\n\nAnbei der derzeitige TOM-Vermerk. Zur Speicherdauer der Logs und der Sicherungskopien prüfe ich gerade die produktiven Einstellungen. Aus der alten Produktfolie würde ich keine verbindlichen Werte übernehmen. Grüße, Sami Arslan','06_TOM_Bestandsaufnahme.pdf'),
        ('13_Fachbereich_2026-10-08.eml','2026-10-08','Juno Nguyễn','personal','Personalordner ist kein Kundenraum','Guten Tag Frau Kranich, unsere zwölf Personalakten liegen im internen Bereich. Die Daten werden monatlich für die Lohnabrechnung vorbereitet. Niemand von unseren Hostingkunden gibt dafür Weisungen. Im alten Blatt habe ich Bewerbungen mit eingetragen, weil die Ordner nebeneinander liegen.\n\nDie vollständige Kundenliste kann der Vertrieb liefern. Die zwei Kontakte im bisherigen Register habe ich aus den greifbaren Verträgen übernommen. Die übrigen Daten wollte ich nicht raten. Mit freundlichen Grüßen, Juno Nguyễn',None),
        ('14_Anbieter_2026-10-09.eml','2026-10-09','Niamh Dorn','anbieter','Supportorte im Angebot','Guten Tag Frau Kranich, der Speicherort Ihrer Daten ändert sich durch unseren Bereitschaftsdienst nicht. Zugriffe können jedoch auch von unseren Spezialisten in Toronto oder Bengaluru erfolgen. Deshalb bitten wir vor einem Test mit echten Inhalten um die abgestimmte Berechtigungsmatrix.\n\nDas Angebot im Anhang nennt die technischen Optionen. Die endgültigen Vertragsanlagen und die beteiligten Gesellschaften werden noch zusammengestellt. Wir haben keinen Produktivzugang erhalten und erwarten auch noch keine Freigabe für reale Kundendaten. Freundliche Grüße, Niamh Dorn','10_Anbieterangebot.pdf')
      ]))
    cases.append(dict(
      slug='vvt-praxis-rosenquell-bamberg',label='Physiopraxis Rosenquell in Bamberg',company='Physiopraxis Rosenquell',address='Rosengärtlein 14, 96050 Bamberg',domain='rosenquell.example',leader='Ottilie Rosenquell',contact='Mina Hartmann',it='Levin Okafor',department='Anselm Fuchs',recipient='Nora Lind, Datenschutzberatung, Quellenhof 6, 96047 Bamberg',email='nora.lind@lind-beratung.example',employees=8,
      profile=[
        'In der Physiopraxis Rosenquell arbeiten die Inhaberin, fünf weitere therapeutische Kräfte und zwei Personen in der Verwaltung. Die Praxis führt laufend Behandlungsdokumentation und Terminplanung. Verordnungen und Abrechnungsangaben werden in der Praxissoftware und teilweise auf Papier bearbeitet.',
        'Ein Anbieter möchte in einem Pilotprojekt kurze Trainingsvideos auswerten und Fortschrittsanzeigen vorschlagen. Geplant sind zunächst 15 teilnehmende Personen. Darunter könnten Minderjährige sein. Ein Test mit echten Videos wurde noch nicht begonnen. Die Inhaberin möchte erst den Verarbeitungsablauf und die notwendigen Entscheidungen klären.',
        'Der Anbieter bezeichnet Videos ohne Namensfeld als anonym. Nach der technischen Beschreibung bleiben Körper, Bewegungsablauf und teilweise Stimme im Bild; der Praxis soll eine Kennung zur Zuordnung dienen. Ob und wie diese Kennung beim Anbieter verwendet wird, ist noch offen.',
        'Bitte ergänzen Sie das Verzeichnis der bestehenden Tätigkeiten und erfassen Sie das geplante Vorhaben getrennt. Stellen Sie nachvollziehbar dar, welche Tatsachen für die weitere Risikoprüfung fehlen und wer sie liefern muss. Eine pauschale Aussage allein anhand der Praxisgröße oder des Wortes Gesundheitsdaten ist nicht beauftragt.'
      ],
      activities=[
        activity('VT-0001','Behandlungsdokumentation und Abrechnung',owner='Ottilie Rosenquell',purpose='Therapeutische Behandlung dokumentieren, Verordnungen verwalten und erbrachte Leistungen abrechnen.',data_subjects='Patientinnen und Patienten, auch Minderjährige; gegebenenfalls Kontaktpersonen.',data_categories='Stamm- und Kontaktdaten, Verordnungen, Befunde, Behandlungseinträge und Abrechnungsdaten.',recipients='Behandelnde Kräfte, Abrechnungsstelle Salbei; weitere Übermittlungen sind zu konkretisieren.',systems='Praxissoftware Rosenblatt, abschließbare Papierablage.',toms='Persönliche Kennungen; Berechtigungsabgrenzung zwischen Empfang und Behandlung noch zu prüfen.',sources='01_Unternehmensprofil_und_Auftrag.docx; 06_TOM_Bestandsaufnahme.pdf'),
        activity('VT-0002','Terminplanung und Erinnerungen',owner='Mina Hartmann',purpose='Termine vereinbaren, Änderungen mitteilen und auf Wunsch an Termine erinnern.',data_subjects='Patientinnen und Patienten sowie anfragende Personen.',data_categories='Name, Kontaktdaten, Terminzeit, behandelnde Person; Freitext kann Behandlungsangaben enthalten.',recipients='Empfang; Versanddienst des Terminprogramms noch zu bestimmen.',transfers='unbekannt',systems='Terminmodul der Praxissoftware und Telefon.',sources='07_Loeschkonzept_Arbeitsstand.pdf; 13_Fachbereich_2026-10-08.eml')
      ],
      change_activity=activity('VT-0003','Geplanter Pilot mit Trainingsvideo und KI-Auswertung',owner='Ottilie Rosenquell',status='geplant',purpose='Geplant: Verlauf einzelner Übungen zeigen und therapeutische Bewertung durch eine Fortschrittsanzeige unterstützen. Kein Beginn mit echten Videos.',data_subjects='Vorgesehen sind 15 freiwillig teilnehmende Personen, möglicherweise Minderjährige.',data_categories='Trainingsvideo mit Körperdarstellung und gegebenenfalls Stimme, Kennung und abgeleitete Bewegungswerte.',recipients='Behandelnde Kräfte; Anbieterzugriffe und weitere Nutzung noch nicht geklärt.',transfers='unbekannt',retention='Anbieter nennt 90 Tage als Einstellung; eine Praxisentscheidung fehlt.',systems='Bewegungsfaden Pilotportal und KI-Modul.',toms='Anbieter wirbt mit Entfernung des Namensfelds. Zuordnung in der Praxis und weitere Identifizierbarkeit sind noch zu prüfen.',sources='08_Vorhaben_und_Aenderungen.docx; 10_Anbieterangebot.pdf'),
      changes='Der neue Eintrag beschreibt nur den noch nicht begonnenen Videopiloten. Der Status geplant hält fest, dass noch kein Produktivbetrieb läuft. Risikoprüfung, Rechtsgrundlagen, Rollen und weitere Maßnahmen sind offen; der Eintrag ist keine Freigabe.',
      avv=[
        'Mina Hartmann hat am 06.10.2026 die Vertragsunterlagen der Praxissoftware und der Abrechnungsstelle bereitgelegt. Zur Software liegt eine unterschriebene Dienstleistungsvereinbarung mit einer Anlage zu technischen Maßnahmen vor. Die Anlage ist vom Februar 2024 und wurde nach einem Versionswechsel noch nicht bestätigt.',
        'Die Abrechnungsstelle erhält die hierfür benötigten Behandlungs- und Kostenträgerangaben. Ihr Vertrag beschreibt eigene organisatorische Aufgaben. Eine pauschale Einordnung sämtlicher Empfänger als gleichartige Dienstleister enthält der Vertragsbestand nicht; diese Einordnung ist im Verzeichnis zu klären.',
        'Zum Videopiloten liegen ein Angebot und ein nicht unterzeichneter Vertragstext vor. Der Anbieter möchte die Auswertungen nach eigener Aussage auch zur Verbesserung seines Modells verwenden. Ob diese Nutzung ausgeschlossen werden kann und welche Rolle er dabei beansprucht, ist im vorgelegten Stand nicht geklärt.'
      ],
      tom=[
        'Levin Okafor hat am 07.10.2026 die Zugänge geprüft. Empfang und Behandlungsräume verwenden persönliche Kennungen. Die Rechte eines ehemaligen Therapeuten wurden deaktiviert. Im Terminfreitext können die Mitarbeitenden derzeit unabhängig vom Aufgabenbereich lesen; die tatsächlich eingetragenen Inhalte wurden nicht vollständig durchgesehen.',
        'Die Notebooks sind verschlüsselt. Papierunterlagen werden abends eingeschlossen. Eine verschlüsselte Sicherung wird erstellt; der jüngste dokumentierte Wiederherstellungstest ist vom 12.06.2026. Eine systematische Prüfung von Versandfehlern bei Erinnerungsnachrichten ist noch nicht dokumentiert.',
        'Für den Videopiloten existiert nur ein Testzugang. Es wurden noch keine Patientenbilder hochgeladen. Ein Entfernen des Namensfelds beseitigt nicht automatisch Körpermerkmale oder Stimmen im Video. Eine verbindliche Beschreibung der Übertragung, Empfänger und Zugriffsmöglichkeiten wurde beim Anbieter angefragt.'
      ],
      retention=[
        'Behandlungsunterlagen, Abrechnungsdaten und Terminkommunikation sind bisher nicht in einem gemeinsamen Löschplan beschrieben. Die Praxis hat die genauen Aufbewahrungsregeln noch nicht in das Verzeichnis übertragen. Das Terminprogramm bietet eine Löschfunktion, die bislang manuell bedient wird.',
        'Absagen und Erinnerungsnachrichten bleiben teilweise im Postfach. Mina Hartmann möchte wissen, ob für diese Nachrichten dieselben Gründe wie für die Behandlungsdokumentation bestehen. Eine allgemeine Antwort für alle Dokumentarten wurde intern noch nicht beschlossen.',
        'Der Videoanbieter nennt 90 Tage als Standardwert und möchte abgeleitete Daten länger behalten. Diese unterschiedlichen Datenbestände und Zwecke wurden noch nicht getrennt beschrieben. Die Praxis hat weder eine Löschdauer noch eine weitere Nutzung für den geplanten Piloten freigegeben.'
      ],
      proposal=[
        'Bewegungsfaden Analyse GmbH, Fadenwiese 9, 91052 Erlangen | Angebot BF-261005 vom 05.10.2026 an Physiopraxis Rosenquell',
        'Wir bieten einen vierwöchigen Piloten für bis zu 15 Teilnehmende an. Die Praxis lädt kurze Übungsvideos mit einer Kennung hoch. Unser Modul erstellt Bewegungswerte und eine vorgeschlagene Fortschrittsanzeige. Die therapeutische Einschätzung verbleibt nach unserem Produktkonzept bei der behandelnden Person.',
        'Videos werden in der Voreinstellung 90 Tage gespeichert. Abgeleitete Daten können zur Modellverbesserung verwendet werden. Bitte teilen Sie uns mit, ob Ihre Praxis diese Option ausschließen möchte. Die Kennung kann in der Praxis mit dem Namen verbunden werden; wir erhalten nach dem bisherigen Konzept kein gesondertes Namensfeld.',
        'Die Broschüre verwendet die Bezeichnung anonym. Eine genauere technische Darstellung der Identifizierbarkeit, der Speicherorte und der Zugriffsmöglichkeiten liefern wir auf Anfrage. Eine abschließende Vereinbarung zur Nutzung von Daten Minderjähriger ist nicht Teil dieses Angebots. Der Pilot ist noch nicht gestartet.'
      ],
      initiative=[
        'Besprechung vom 06.10.2026 mit Ottilie Rosenquell, Anselm Fuchs und Mina Hartmann. Herr Fuchs möchte zunächst nur kurze Videos einer festgelegten Übung aufnehmen. Frau Rosenquell möchte vor der Auswahl von Teilnehmenden wissen, welche Daten das System tatsächlich benötigt.',
        'Die Anbieterfolie spricht von anonymen Videos. Frau Hartmann weist darauf hin, dass Gesichter in manchen Einstellungen zu sehen wären und Stimmen mit aufgenommen werden könnten. Eine Zuordnungsliste soll in der Praxis geführt werden. Die endgültige Kameraperspektive steht nicht fest.',
        'Herr Fuchs hält eine Fortschrittsanzeige für hilfreich, möchte aber keine Therapieentscheidung ungeprüft vom System übernehmen. Die Inhaberin hat eine Prüfung der möglichen Folgen von Fehlbewertungen, der Datenweiterverwendung und der Beteiligung Minderjähriger beauftragt. Bis dahin soll ausschließlich mit erfundenen oder nachweisbar personenfreien Testdaten gearbeitet werden.'
      ],
      requests=[
        'bitte erläutern Sie mir bis zum 19.10.2026 den vollständigen Ablauf des geplanten Videopiloten. Ich benötige die aufgenommenen Bild- und Tonbestandteile, die Kennung, die in der Praxis vorhandene Zuordnung und die erzeugten Auswertungen. Bitte beschreiben Sie auch, welche Bedeutung eine Fortschrittsanzeige für die Behandlung tatsächlich erhalten soll und wie fehlerhafte Ergebnisse bemerkt werden.',
        'Vom Anbieter benötige ich die tatsächlichen Speicher- und Supportorte, sämtliche Empfänger, die vorgesehenen Zugriffsmöglichkeiten und eine getrennte Beschreibung der Modellverbesserung. Bitte lassen Sie ausdrücklich erläutern, auf welcher technischen Grundlage der Anbieter die Videos als anonym bezeichnet. Die fehlende Namensspalte allein beantwortet diese Frage nicht.',
        'Für die bestehenden Verfahren ergänzen Sie bitte die Zuständigkeiten, Empfänger und dokumentierten Aufbewahrungsregeln nach Datenarten. Zum Pilotprojekt bleiben Risikoprüfung und weitere Entscheidung offen. Bitte wählen Sie bis dahin keine echten Patientenvideos aus und holen Sie keine bloß formale Zustimmung ein, bevor der konkrete Ablauf geklärt ist.'
      ],
      mails=[
        ('11_GF_2026-10-06.eml','2026-10-06','Ottilie Rosenquell','praxisleitung','Kleiner Pilot erst nach Klärung','Guten Tag Frau Lind, wir möchten den Videopiloten verstehen, bevor wir ihn beginnen. Unsere Praxis ist klein, aber wir bearbeiten natürlich jeden Tag Behandlungsdaten. Bitte prüfen Sie die bestehenden Abläufe und das neue Vorhaben jeweils konkret.\n\nHerr Fuchs hat bisher nur die Broschüre gesehen. Es gibt noch keine Teilnehmendenliste und keine hochgeladenen Videos. Eine allgemeine Unterschrift unserer Patientinnen und Patienten möchte ich nicht einholen, solange wir den Ablauf selbst noch nicht erklären können. Mit freundlichen Grüßen, Ottilie Rosenquell',None),
        ('12_IT_2026-10-07.eml','2026-10-07','Levin Okafor','it','Was im Video sichtbar bleibt','Guten Tag Frau Lind, das Portal entfernt nach dem Angebot nur ein Namensfeld. Das Video selbst kann weiterhin Körper, Gesicht und Stimme enthalten. Außerdem bleibt die Kennung mit einer Liste in der Praxis verbunden. Ob der Anbieter diese Kennung auch für andere Zwecke benutzt, weiß ich nicht.\n\nDie Bestandsaufnahme der bisherigen IT liegt bei. Für den neuen Dienst gibt es lediglich einen Testzugang. Bitte übernehmen Sie den Broschürentext anonym nicht ungeprüft als Beschreibung unseres Vorhabens. Grüße, Levin Okafor','06_TOM_Bestandsaufnahme.pdf'),
        ('13_Fachbereich_2026-10-08.eml','2026-10-08','Anselm Fuchs','therapie','Fortschrittsanzeige und Terminnotizen','Guten Tag Frau Lind, ich möchte die Anzeige nur zusätzlich zu meiner eigenen Beobachtung nutzen. Ein schlechter Wert darf nicht allein darüber entscheiden, ob jemand eine weitere Behandlung erhält. Unter den Interessierten könnten zwei Jugendliche sein, aber noch niemand wurde angesprochen.\n\nIm Terminmodul steht manchmal der Grund einer Verschiebung, etwa Schmerzen nach einer Operation. Die Übersicht ist also nicht immer nur ein Kalender mit Namen. Frau Hartmann kann die Zugriffsrechte zeigen. Mit freundlichen Grüßen, Anselm Fuchs',None),
        ('14_Anbieter_2026-10-09.eml','2026-10-09','Alma Stern','anbieter','Rückfragen zum Bewegungsfaden Piloten','Guten Tag Frau Lind, die Bezeichnung anonym in unserer Broschüre bezieht sich auf das fehlende Namensfeld. Eine weitergehende technische Beschreibung reichen wir nach. Eine Kennung ist für die Verlaufsansicht erforderlich. Der Zugriff auf die Zuordnungsliste in Ihrer Praxis ist nach dem bisherigen Konzept nicht vorgesehen.\n\nDas Angebot liegt bei. Wir klären intern, ob die Modellverbesserung vollständig ausgeschlossen werden kann und welche Speicher- und Supportorte für Ihren Pilot gelten. Bitte laden Sie bis zu dieser Klärung keine echten Videos hoch. Freundliche Grüße, Alma Stern','10_Anbieterangebot.pdf')
      ]))
    for c in cases:
        c['provider_domain']={'vvt-handwerk-finkenbeil-erfurt':'spurspatz.example','vvt-cloudservice-wolkengarn-berlin':'nachtfaden.example','vvt-praxis-rosenquell-bamberg':'bewegungsfaden.example'}[c['slug']]
        c['provider_address']={'vvt-handwerk-finkenbeil-erfurt':'Spurspatz Mobilitätssoftware GmbH, Antennenhain 5, 04129 Leipzig','vvt-cloudservice-wolkengarn-berlin':'Nachtfaden Support Ltd., Supportkoordination Dublin, Kontakt über support@nachtfaden.example','vvt-praxis-rosenquell-bamberg':'Bewegungsfaden Analyse GmbH, Fadenwiese 9, 91052 Erlangen'}[c['slug']]
        base=dict(schema_version=1,register_id=str(uuid.uuid5(uuid.NAMESPACE_URL,c['slug'])),organization=dict(name=c['company'],contact=f"{c['address']}; {c['contact']}; buero@{c['domain']}",representative='',dpo=''),revision=1,updated_at='2026-10-07T16:00:00+02:00',activities=c['activities'],history=[])
        after=copy.deepcopy(base);after['revision']=2;after['updated_at']='2026-10-10T09:00:00+02:00'
        if c['change_activity']:after['activities'].append(c['change_activity'])
        else:
            a=after['activities'][0];a['activity_revision']=2;a['transfers']='Bisher Rechenzentren Deutschland und Finnland. Geplanter Nachtsupport nennt Toronto und Bengaluru. Zugriffsrechte, Gesellschaften und dokumentierte Grundlage sind offen; kein Produktivzugang eingerichtet.';a['sources']+='; 14_Anbieter_2026-10-09.eml'
        after['history']=[dict(revision=2,at='2026-10-10T09:00:00+02:00',actor=c['contact'],reason=c['changes'],operation='Entwurfsstand ergänzt',changed_ids=['VT-0003'] if c['change_activity'] else ['VT-0001'])]
        c['before']=base;c['after']=after
    return cases

def build(cases):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.oxml.ns import qn
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
    fonts=Path('/System/Library/Fonts/Supplemental')
    pdfmetrics.registerFont(TTFont('VvtSerif',str(fonts/'Times New Roman.ttf')));pdfmetrics.registerFont(TTFont('VvtBold',str(fonts/'Times New Roman Bold.ttf')))
    ps=ParagraphStyle('Body',fontName='VvtSerif',fontSize=11,leading=15,spaceAfter=10)
    ts=ParagraphStyle('Title',parent=ps,fontName='VvtBold',fontSize=17,leading=21,spaceAfter=16)
    for c in cases:
        folder=ROOT/'testakten'/c['slug'];folder.mkdir(parents=True,exist_ok=True)
        def docx(name,title,paras):
            d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(22);s.left_margin=s.right_margin=Mm(24)
            for n in ['Normal','Title','Heading 1']:
                st=d.styles[n];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
                for attr in list(st.element.rPr.rFonts.attrib):
                    if attr.endswith('Theme') or attr.endswith('theme'):del st.element.rPr.rFonts.attrib[attr]
                for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
            d.styles['Normal'].paragraph_format.space_after=Pt(9);d.styles['Normal'].paragraph_format.line_spacing=1.08
            d.styles['Title'].font.size=Pt(17);d.styles['Title'].font.bold=True
            d.add_paragraph(title,'Title')
            for text in paras:d.add_paragraph(text)
            d.core_properties.author=c['company'];d.core_properties.title=title;d.save(folder/name)
        def pdf(name,title,paras):
            story=[Paragraph(escape(title),ts),Paragraph(escape(c['company']+' | '+c['address']),ps),Spacer(1,8)]
            for text in paras:story.append(Paragraph(escape(text),ps))
            def page(can,doc):can.setFont('VvtSerif',9);can.drawString(54,28,name);can.drawRightString(A4[0]-54,28,str(doc.page))
            SimpleDocTemplate(str(folder/name),pagesize=A4,leftMargin=54,rightMargin=54,topMargin=50,bottomMargin=48,title=title,author=c['company']).build(story,onFirstPage=page,onLaterPages=page)
        docx('01_Unternehmensprofil_und_Auftrag.docx','Unternehmensprofil und Arbeitsauftrag',[c['company']+'\n'+c['address']+'\nBearbeitungsstand 10.10.2026']+c['profile'])
        docx('08_Vorhaben_und_Aenderungen.docx','Geplante Änderung der Verarbeitung',c['initiative'])
        docx('09_Nachforderung_Entwurf.docx','Unterlagen zum Verarbeitungsverzeichnis',[c['recipient']+'\n'+c['email'],c['company']+'\nAn die Geschäftsführung\n'+c['address']+'\n10.10.2026','Sehr geehrte '+('Frau '+c['leader'].split()[-1] if c['slug']!='vvt-cloudservice-wolkengarn-berlin' else 'Damen und Herren')+',']+c['requests']+['Mit freundlichen Grüßen\n'+c['recipient'].split(',')[0]+'\nDatenschutzberatung'])
        pdf('05_AVV_Vertragsbestand.pdf','Vermerk zum Vertragsbestand',c['avv'])
        pdf('06_TOM_Bestandsaufnahme.pdf','Bestandsaufnahme der technischen Maßnahmen',c['tom'])
        pdf('07_Loeschkonzept_Arbeitsstand.pdf','Arbeitsstand zur Aufbewahrung und Löschung',c['retention'])
        pdf('10_Anbieterangebot.pdf','Angebot für das geplante Vorhaben',c['proposal'])
        (folder/'02_Registerbestand.json').write_text(json.dumps(c['before'],ensure_ascii=False,indent=2)+'\n')
        (folder/'03_Register_Aenderungsstand.json').write_text(json.dumps(c['after'],ensure_ascii=False,indent=2)+'\n')
        for n,(name,date,sender,mailbox,subject,body,attachment) in enumerate(c['mails']):
            sender_domain=c['provider_domain'] if mailbox=='anbieter' else c['domain']
            msg=EmailMessage(policy=SMTP);msg['From']=f'{sender} <{mailbox}@{sender_domain}>';msg['To']=c['email'];msg['Subject']=subject
            msg['Date']=format_datetime(dt.datetime.fromisoformat(date+'T10:15:00').replace(tzinfo=ZoneInfo('Europe/Berlin')));msg['Message-ID']=f'<{c["slug"]}-{n}-{date}@{c["domain"]}>'
            msg.set_content(body+'\n\n'+sender+'\n'+(c['provider_address'] if mailbox=='anbieter' else c['address']))
            if attachment:msg.add_attachment((folder/attachment).read_bytes(),maintype='application',subtype='pdf',filename=attachment)
            (folder/name).write_bytes(msg.as_bytes())
        readme=f'# Testakte {c["label"]}\n\nDiese Akte dient zur Aufnahme, Prüfung und Pflege eines Verzeichnisses von Verarbeitungstätigkeiten. Ausgangspunkt sind das Unternehmensprofil und der Arbeitsauftrag. Fachbereiche liefern teils vollständige, teils widersprüchliche Angaben; geplante Änderungen sind noch nicht freigegeben.\n\n## 1. Bearbeitung\n\nÖffnen Sie die Einzelunterlagen und die bearbeitbare Excel-Datei. Vergleichen Sie den älteren Registerbestand mit dem jüngeren Änderungsstand. Der jüngere Stand enthält zusätzliche Angaben aus den Aktenstücken, aber keine abschließende rechtliche Bewertung. Die offenen Felder bleiben sichtbar. Die Datei `09_Nachforderung_Entwurf.docx` ist ein ausformulierter vorhandener Briefentwurf, den Sie mit dem ermittelten Nachforderungsbedarf abgleichen können.\n\nDie zwei JSON-Fassungen enthalten Tätigkeiten mit stabilen Kennungen. Der Status `geplant` kennzeichnet noch nicht begonnene Vorhaben. Aus Eintrag, Screening oder Ampel folgt keine Startfreigabe.\n\n## 2. Fiktive Unterlagen\n\nAlle Unternehmen, Personen, Adressen und Vorgänge sind erfunden. Die reservierten `.example`-Adressen sind keine Versandziele. Die Akte enthält keine echten Patienten- oder Beschäftigtendaten und keine Musterlösung.\n\n<!-- reserved-example-contacts -->\n\n<!-- BEGIN gesamt-pdf-section (autogen) -->\n## 3. Downloads\n\n{WARNING}\n\n| Format | Download |\n| --- | --- |\n| Gesamt-PDF | [Lesefassung](gesamt-pdf/{c["slug"]}_gesamt.pdf) |\n| Originale | [Akten-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/verarbeitungsverzeichnis-v1.0.0/testakte-{c["slug"]}.zip) |\n| Einzel-PDFs | [Einzel-PDF-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/verarbeitungsverzeichnis-v1.0.0/testakte-{c["slug"]}-einzelpdfs.zip) |\n\n<!-- END gesamt-pdf-section (autogen) -->\n'
        (folder/'README.md').write_text(readme)
        (folder/'rubric.yaml').write_text(f'name: {c["slug"]}\nplugin: verarbeitungsverzeichnis\ndescription: "{c["label"]}: Verzeichnisbestand und offene Änderungen mit nativen Unterlagen."\nchecks:\n  - id: unterlagen\n    check_type: working_file_count\n    description: "Mindestens zwölf eigenständige Originalunterlagen liegen vor."\n    min: 12\n  - id: rollen\n    check_type: human_review\n    description: "Rolle und tatsächliche Verarbeitung werden aus dem Sachverhalt begründet und nicht pauschal übernommen."\n  - id: aenderung\n    check_type: human_review\n    description: "Einträge behalten stabile Kennungen; Änderungen ersetzen keine fachliche Freigabe."\n  - id: luecken\n    check_type: human_review\n    description: "Fehlende Angaben werden konkret nachgefordert und nicht erfunden."\n')
    print('Drei Akten erstellt; der VVT-Helfer exportiert anschließend die Excel-Dateien.')

def main():
    p=argparse.ArgumentParser();p.add_argument('--data-only',action='store_true');p.add_argument('--exports-only',action='store_true');args=p.parse_args();cases=dataset();DATA.parent.mkdir(parents=True,exist_ok=True);DATA.write_text(json.dumps(cases,ensure_ascii=False,indent=2)+'\n')
    if not args.data_only and not args.exports_only:build(cases)
    if not args.data_only:
        import importlib.util
        spec=importlib.util.spec_from_file_location('vvt',ROOT/'verarbeitungsverzeichnis/scripts/vvt.py');vvt=importlib.util.module_from_spec(spec);spec.loader.exec_module(vvt)
        for c in cases:
            vvt.validate(c['before']);vvt.validate(c['after'])
            vvt.export_xlsx(c['after'],ROOT/'testakten'/c['slug']/'04_Verzeichnis.xlsx')
        template=ROOT/'verarbeitungsverzeichnis/templates/startregister';template.mkdir(parents=True,exist_ok=True)
        start=dict(schema_version=1,register_id=str(uuid.uuid5(uuid.NAMESPACE_URL,'verarbeitungsverzeichnis/startregister/v1')),organization=dict(name='[Unternehmen ergänzen]',contact='',representative='',dpo=''),revision=0,updated_at='2026-10-10T09:00:00+02:00',activities=[],history=[])
        vvt.validate(start);(template/'verarbeitungsverzeichnis.json').write_bytes(vvt.canonical(start))
        for fmt in ['xlsx','docx','xml','html']:getattr(vvt,'export_'+fmt)(start,template/('verarbeitungsverzeichnis.'+fmt))
        print('Drei importierbare Excelregister und Startregister in fünf Formaten exportiert.')

if __name__=='__main__':main()
