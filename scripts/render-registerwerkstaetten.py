#!/usr/bin/env python3
"""Lesefassungen der vier Registerwerkstätten: TNR 11, native Office-Ausgabe, QA-PNGs."""
from pathlib import Path
import argparse,hashlib,importlib.util,json,re,subprocess,sys,io,zipfile
from datetime import datetime,timezone
from urllib.parse import quote
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader,PdfWriter
from pypdf.generic import ArrayObject,ByteStringObject,NameObject
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('workshop_base',ROOT/'scripts/render-startup-gruender-werkstatt.py');base=importlib.util.module_from_spec(spec);spec.loader.exec_module(base)
DATE=datetime(2026,10,1,9,tzinfo=timezone.utc)
def clean(s):return re.sub(r'\*\*([^*]+)\*\*|`([^`]+)`',lambda m:m[1] or m[2],s)
def build(text,target,title):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
 s.left_margin=s.right_margin=Cm(2.4);s.top_margin=s.bottom_margin=Cm(2.2);s.header_distance=s.footer_distance=Cm(1.2)
 for name,size in [('Normal',11),('Title',20),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
  st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.line_spacing=1.5 if name=='Normal' else 1.15;st.paragraph_format.space_after=Pt(8)
  st.paragraph_format.widow_control=True
  if name!='Normal':st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(14)
 for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 for node in d.styles.element.iter(qn('w:rFonts')):
  for key in list(node.attrib):
   if 'theme' in key.lower():del node.attrib[key]
  node.set(qn('w:ascii'),'Times New Roman');node.set(qn('w:hAnsi'),'Times New Roman')
 d.core_properties.title=title;d.core_properties.author='Registerwerkstätten';d.core_properties.created=d.core_properties.modified=DATE
 lines=text.splitlines();i=0;headings=0
 while i<len(lines):
  line=lines[i].strip()
  if not line or line.startswith('<!--'):i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].lstrip().startswith('|'):
    cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[: -]+',x) for x in cells):rows.append(cells)
    i+=1
   t=d.add_table(rows=1,cols=len(rows[0]));t.style='Table Grid'
   for n,row in enumerate(rows):
    cells=t.rows[0].cells if n==0 else t.add_row().cells
    for cell,value in zip(cells,row):
     base.inline(cell.paragraphs[0],clean(value));cell.paragraphs[0].paragraph_format.line_spacing=1.1
     for run in cell.paragraphs[0].runs:run.font.size=Pt(10)
   repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
   for row in t.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   continue
  m=re.match(r'^(#{1,4})\s+(.*)',line)
  if m:
   level=len(m[1]);style='Title' if level==1 else f'Heading {min(level-1,3)}';p=d.add_paragraph(style=style);base.inline(p,clean(m[2]));headings+=int(level>1);i+=1;continue
  if line.startswith('- '):
   p=d.add_paragraph();p.paragraph_format.left_indent=Cm(.35);base.inline(p,'• '+clean(line[2:]));i+=1;continue
  if line.startswith('```'):raise ValueError('Codeblock muss gezielt gerendert werden: '+str(target))
  parts=[line];i+=1
  while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||- |<!--)',lines[i]):parts.append(lines[i].strip());i+=1
  p=d.add_paragraph();base.inline(p,clean(' '.join(parts)))
 footer=s.footer.paragraphs[0];footer.alignment=2;r=footer.add_run(title+' · ');r.font.name='Times New Roman';r.font.size=Pt(9)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 memory=io.BytesIO();d.save(memory)
 with zipfile.ZipFile(memory) as src,zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as dst:
  for name in sorted(src.namelist()):
   info=zipfile.ZipInfo(name,(2026,10,1,9,0,0));info.compress_type=zipfile.ZIP_DEFLATED;dst.writestr(info,src.read(name))
 return headings

def main():
 p=argparse.ArgumentParser();p.add_argument('plugins',nargs='+');p.add_argument('--qa-dir',type=Path,required=True);p.add_argument('--renderer',type=Path,required=True);a=p.parse_args()
 runtime=Path(sys.executable).resolve().parents[3]
 for slug in a.plugins:
  source=ROOT/slug/f'{slug}-werkstatt.md';text=source.read_text();title=source.read_text().splitlines()[0].lstrip('# ')
  def absolute_link(match):
   label,url=match.groups()
   if re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:',url) or url.startswith('#'):return match.group(0)
   path=(source.parent/url).resolve().relative_to(ROOT)
   return '['+label+'](https://github.com/Klotzkette/claude-fuer-deutsches-recht/blob/main/'+quote(path.as_posix(),safe='/')+')'
  text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',absolute_link,text)
  out=ROOT/slug/'assets';out.mkdir(exist_ok=True);qa=a.qa_dir/slug;qa.mkdir(parents=True,exist_ok=True);render=qa/'render';render.mkdir(exist_ok=True)
  docx=out/f'{slug}-werkstatt.docx';headings=build(text,docx,title);digest=hashlib.sha256(source.read_bytes()).hexdigest()
  for obsolete in render.glob('page-*.png'):obsolete.unlink()
  env=base.renderer_environment(qa,runtime)
  result=subprocess.run([sys.executable,str(a.renderer),str(docx),'--output_dir',str(render),'--emit_pdf','--dpi','100'],env=env,capture_output=True,text=True,timeout=300)
  (qa/'render.log').write_text(result.stdout+result.stderr)
  if result.returncode:raise RuntimeError((result.stdout+result.stderr)[-4000:])
  raw=render/f'{slug}-werkstatt.pdf';reader=PdfReader(raw);writer=PdfWriter();writer.clone_document_from_reader(reader)
  writer.add_metadata({'/Title':title,'/Author':'Registerwerkstätten','/CreationDate':'D:20261001090000Z','/ModDate':'D:20261001090000Z'})
  if '/Metadata' in writer._root_object:del writer._root_object[NameObject('/Metadata')]
  token=bytes.fromhex(digest[:32]);writer._ID=ArrayObject([ByteStringObject(token),ByteStringObject(token)])
  target=out/f'{slug}-werkstatt.pdf'
  with target.open('wb') as f:writer.write(f)
  info={'plugin':slug,'source_sha256':digest,'source_words':len(text.split()),'headings':headings,'pages':len(reader.pages),'font':'Times New Roman','body_points':11,'line_spacing':1.5,'manual_page_breaks':0,'visual_review':'pending'}
  (qa/'render-result.json').write_text(json.dumps(info,ensure_ascii=False,indent=2)+'\n');print(json.dumps(info),flush=True)
if __name__=='__main__':main()
