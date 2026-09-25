#!/usr/bin/env python3
"""Fallbezogene Regression ohne zentrale Änderungen. Autor: Klotzkette."""
from __future__ import annotations
import argparse
from collections import Counter
from datetime import datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request
from unittest.mock import patch
import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageStat
from pypdf import PdfReader
from testakte_disclaimer import NOTICE_BYTES
from testakte_file_filter import include_in_working_dump

HERE=Path(__file__).resolve().parent
def load(name,filename):
    spec=importlib.util.spec_from_file_location(name,HERE/filename);m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m
B=load("einbeck_case_test","build-bauwirtschaft-einbeck-akte.py")
V=load("einbeck_validator","validate-testakten-dokumentqualitaet.py")
SOURCES=[
 ("HOAI Anlage 10", "https://www.gesetze-im-internet.de/hoai_2013/anlage_10.html", "Objektbetreuung"),
 ("HOAI Paragraf 34", "https://www.gesetze-im-internet.de/hoai_2013/__34.html", "Leistungsphasen"),
 ("NBauO Paragraf 70", "https://voris.wolterskluwer-online.de/browse/document/dual/87757479-5739-3a13-a220-fbbcc613b1e5/c3e55212-5780-362d-9beb-c20b1d5fac1e", "Baugenehmigung"),
 ("Einbeck Bauaufsicht", "https://www.einbeck.de/buergerservice/dienstleistungen/bauaufsicht-900000569-0.html", "Bauaufsicht"),
 ("BGB Paragraf 634a", "https://www.gesetze-im-internet.de/bgb/__634a.html", "Verj"),
 ("Niedersachsen Verwaltungsgerichtsprozess", "https://justizportal.niedersachsen.de/startseite/gerichte_und_staatsanwaltschaften/verwaltungsgerichtsbarkeit/verwaltungsgerichtsprozess/", "Widerspruch"),
]

def sources(qa):
    dest=qa/"quellen";dest.mkdir(parents=True,exist_ok=True);result=[]
    for i,(title,url,needle) in enumerate(SOURCES):
        if shutil.which("curl"):
            # Der Systemclient nutzt den lokalen Zertifikatsspeicher; TLS bleibt geprüft.
            response=subprocess.run(["curl","--fail","--location","--silent","--show-error","--max-time","40",
                                     "--user-agent","Mozilla/5.0","--write-out","\n%{http_code}",url],check=True,capture_output=True)
            raw,code=response.stdout.rsplit(b"\n",1);status=int(code)
        else:
            req=urllib.request.Request(url,headers={"User-Agent":"Mozilla/5.0"})
            with urllib.request.urlopen(req,timeout=40) as response:raw=response.read();status=response.status
        assert status==200 and needle.lower().encode() in raw.lower(),title
        (dest/f"quelle-{i+1}.html").write_bytes(raw)
        result.append(dict(title=title,url=url,status=status,bytes=len(raw),sha256=hashlib.sha256(raw).hexdigest(),retrieved=datetime.now().astimezone().isoformat()))
    (dest/"abruf.json").write_text(json.dumps(result,ensure_ascii=False,indent=2));return result

def native_probes(qa):
    P=load("einbeck_package_test","build-bauwirtschaft-einbeck-pakete.py")
    P.QA=qa
    probes=[(9,"Kosten","B6",217,"E16",337580),(19,"Termine","C6",37,"D6",datetime(2022,6,7)),
            (24,"LV","B6",2,"E19",108285),(37,"Rechnung","B6",2,"E24",123153.8),
            (9,"Kosten","B6",0,"E16",263800),(9,"Kosten","B6",None,"E16","#N/A"),
            (19,"Termine","C6",0,"D6","#N/A")]
    result=[]
    for i,(n,s,c,value,out,expected) in enumerate(probes):
        dest=P.QA/"mutationen"/str(i);dest.mkdir(parents=True,exist_ok=True)
        path=dest/B.FILES[n];shutil.copyfile(B.CASE/B.FILES[n],path)
        P.patch_xlsx(path,dict(sheet=1,cell=c,value=value))
        calculated=P.native_roundtrip([path],f"mutation-{i}")[0]
        wb=load_workbook(calculated,data_only=True);actual=wb[s][out].value;wb.close()
        if isinstance(expected,(int,float)):assert abs(actual-expected)<.001,(n,c,actual,expected)
        else:assert actual==expected,(n,c,actual,expected)
        result.append(dict(file=B.FILES[n],input=f"{s}!{c}",value=value,output=out,actual=str(actual),expected=str(expected)))
    (P.QA/"native-aenderungspruefung.json").write_text(json.dumps(result,ensure_ascii=False,indent=2));return result

def main():
    parser=argparse.ArgumentParser();parser.add_argument("--native",action="store_true");parser.add_argument("--fetch-sources",action="store_true")
    parser.add_argument("--assets",type=Path,help="Optionale Paketprüfung nach dem Release-Build; nie im Standardlauf vorausgesetzt.")
    args=parser.parse_args()
    temporary=tempfile.TemporaryDirectory(prefix="einbeck-quellenpruefung-")
    qa=(args.assets/"qa-einbeck") if args.assets else Path(temporary.name)
    qa.mkdir(parents=True,exist_ok=True)
    with patch.dict(sys.modules,{"pymupdf":None,"fitz":None}):
        load("einbeck_package_dependency_test","build-bauwirtschaft-einbeck-pakete.py")
    original=[p for p in B.CASE.iterdir() if include_in_working_dump(p,B.CASE)]
    assert len(original)==44,len(original)
    assert {p.name for p in original}==set(B.FILES.values())
    assert all(p.suffix in {".docx",".pdf",".xlsx",".eml",".png",".txt",".csv"} for p in original)
    assert sorted(n for ids in B.PHASES.values() for n in ids)==list(range(1,45))
    assert set(B.PHASES)==set(range(1,10)) and all(len(ids)>=2 for ids in B.PHASES.values())
    assert len(B.SLUG)<=64
    B.define_documents()
    texts={}
    for n,name in B.FILES.items():
        path=B.CASE/name;date=datetime.fromisoformat(name[3:13]);assert datetime(2020,1,1)<=date<=datetime(2026,9,25)
        if path.suffix in V.EXPORT_TEXT_EXTS:texts[n]=V.export_text(path)
        if path.suffix==".eml":
            msg=BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            for h in ("From","To","Date","Subject","Message-ID","Return-Path","Received","MIME-Version"):assert msg[h],(name,h)
            attachments=list(msg.iter_attachments());expected=B.DATA[n]["attachments"]
            assert len(attachments)==len(expected)
            for part,num in zip(attachments,expected):
                assert part.get_filename()==B.FILES[num]
                assert hashlib.sha256(part.get_payload(decode=True)).digest()==hashlib.sha256((B.CASE/B.FILES[num]).read_bytes()).digest()
                assert B.FILES[num][3:13]<=name[3:13],(name,part.get_filename())
        elif path.suffix==".docx":
            doc=Document(path);assert doc.core_properties.author=="Klotzkette";assert doc.core_properties.last_modified_by=="Klotzkette"
        elif path.suffix==".pdf":assert PdfReader(path).metadata.author=="Klotzkette"
        elif path.suffix==".xlsx":
            w=load_workbook(path,data_only=False);cache=load_workbook(path,data_only=True)
            assert w.properties.creator=="Klotzkette"
            for s in w:
                for row in s:
                    for c in row:
                        if c.data_type=="f":assert cache[s.title][c.coordinate].value is not None,(name,c.coordinate)
                        assert cache[s.title][c.coordinate].data_type!="e",(name,c.coordinate)
            w.close();cache.close()
        elif path.suffix==".png":
            im=Image.open(path);assert min(im.size)>=1000
            assert max(ImageStat.Stat(im.convert("RGB")).stddev)>20
    forbidden=re.compile(r"\b(?:ChatGPT|Claude|Codex|OpenAI)\b|\u00a7",re.I)
    for n,t in texts.items():
        assert not forbidden.search(t),(n,forbidden.search(t).group())
        assert not re.search(r"(?m)^\s*(?:[IVX]+\.|[A-Z]\.)\s",t),n
        assert "Diese Testakte wurde" not in t and "This test case file" not in t,n
    rescue_pages = PdfReader(B.CASE / B.FILES[16]).pages
    assert len(rescue_pages) == 2
    assert "3 Betrieb und technische Ausstattung" in rescue_pages[-1].extract_text()
    assert "Mit freundlichen Grüßen" in rescue_pages[-1].extract_text()
    assert "3.706,00" in texts[34] and "1.240,00" in texts[34]
    assert "21,45" in texts[35] and "4.410,14" in texts[34]
    assert "18.03.2029" in texts[44] and "02.04.2029" in texts[44]
    assert "Masseprozente" in texts[43] and "Horizontalsperre" in texts[43]
    assert "Keine gesonderte Erklärung" in texts[44]
    amounts={9:("Kosten",{"E16":337240,"F16":401315.6,"F20":28684.4}),24:("LV",{"E19":99785,"E21":118744.15}),
             37:("Rechnung",{"D24":115393.8,"E24":114153.8,"D30":54018.62,"E30":52543.02,"F30":1475.6})}
    for n,(sheet,cells) in amounts.items():
        w=load_workbook(B.CASE/B.FILES[n],data_only=True)
        for cell,value in cells.items():assert abs(w[sheet][cell].value-value)<.001,(n,cell,w[sheet][cell].value,value)
        w.close()
    actual=Decimal("0")
    for qty,ep in zip(B.FINAL_QTY,B.LEINE_EP):actual+=Decimal(str(qty))*Decimal(str(ep))
    assert actual==Decimal("110447.80")
    # Vollständige bestehende Validatorlogik, ausschließlich auf diesen Fall begrenzt.
    scope=qa/"validator";testakten=scope/"testakten";testakten.mkdir(parents=True,exist_ok=True)
    case_link=testakten/B.SLUG
    if not case_link.exists():case_link.symlink_to(B.CASE,target_is_directory=True)
    V.REPO=scope;V.TESTAKTEN=testakten
    assert V.main()==0,"Bestehender Dokumentqualitätsvalidator"
    filtered={p.name for p in B.CASE.iterdir() if include_in_working_dump(p,B.CASE)}
    assert filtered==set(B.FILES.values()),"Exportfilter verliert Dokumente"
    ignored=subprocess.run(["git","check-ignore","--stdin"],input="\n".join(str(p.relative_to(B.ROOT)) for p in original)+"\n",cwd=B.ROOT,capture_output=True,text=True)
    assert ignored.returncode==1 and not ignored.stdout,ignored.stdout
    pdf_pages=None
    if args.assets:
        individual_pages=0
        for variant in ["", "-einzelpdfs"]:
            filename=f"testakte-{B.SLUG}{variant}.zip"
            candidates=[args.assets/filename,args.assets/"testakten-einzelpdfs"/filename,args.assets/"testakten"/filename]
            p=next((p for p in candidates if p.is_file()),None)
            assert p is not None,f"Paket fehlt: {filename}"
            with zipfile.ZipFile(p) as z:
                names=z.namelist();assert len(set(names))==len(names)
                assert all("/" not in name and not name.endswith(".md") for name in names)
                assert z.read("README.txt").startswith(NOTICE_BYTES)
                if variant:
                    assert set(names)=={Path(f).stem+".pdf" for f in B.FILES.values()}|{"README.txt"}
                else:
                    extra=set(names)-set(B.FILES.values())-{"README.txt"}
                    assert extra.issubset({B.SLUG+"_gesamt.pdf","gesamt-pdf__"+B.SLUG+"_gesamt.pdf"}),extra
                    assert set(B.FILES.values()).issubset(names)
                    for name in B.FILES.values():assert z.read(name)==(B.CASE/name).read_bytes(),name
                for name in names:
                    if name.endswith(".pdf"):
                        pdf=PdfReader(io.BytesIO(z.read(name)))
                        text="\n".join(p.extract_text() for p in pdf.pages)
                        assert "Diese Testakte wurde" not in text and "This test case file" not in text
                        if variant:
                            individual_pages+=len(pdf.pages)
                            assert text.strip() or name.startswith(("22_","32_","42_")),name
        candidates=[args.assets/(B.SLUG+"_gesamt.pdf"),args.assets/"gesamt-pdf"/(B.SLUG+"_gesamt.pdf"),B.CASE/"gesamt-pdf"/(B.SLUG+"_gesamt.pdf")]
        total=next((p for p in candidates if p.is_file()),None);assert total is not None,"Gesamt-PDF fehlt"
        reader=PdfReader(total);pdf_pages=len(reader.pages)
        assert pdf_pages>=individual_pages,"Gesamt-PDF hat weniger Seiten als die Einzelunterlagen"
        totaltext="\n".join(p.extract_text() for p in reader.pages)
        assert "Diese Testakte wurde" not in totaltext and "This test case file" not in totaltext
        for needle in ["Baugenehmigung", "LB-240301", "Gerätebereitschaft", "18.03.2029", "02.04.2029"]:assert needle in totaltext,needle
    for n in (4,7,12,13,20,21):
        text=PdfReader(B.CASE/B.FILES[n]).pages[0].extract_text();assert "Maßstab" in text or "Detail 1:10" in text or n==13
    source_result=sources(qa) if args.fetch_sources else []
    probes=native_probes(qa) if args.native else []
    report=dict(originals=44,formats=dict(Counter(p.suffix for p in original)),phases=B.PHASES,pdf_pages=pdf_pages,
                native_probes=probes,sources=source_result,existing_validator="unverändert fallisoliert bestanden",git_ignore="keine Originaldatei ignoriert",result="bestanden")
    (qa/"testbericht.json").write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False,default=str))
    temporary.cleanup()

if __name__=="__main__":main()
