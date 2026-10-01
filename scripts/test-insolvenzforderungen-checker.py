#!/usr/bin/env python3
"""Paket, Fallbelege und Prüfgrenzen; keine behaupteten Modellläufe."""
import csv
from datetime import date
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
import hashlib
import json
from pathlib import Path
import re
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from pypdf import PdfReader
from openpyxl import load_workbook
from prompt_limits import mini_within_limits
from prompt_profiles import validate_files
from testakte_file_filter import include_in_working_dump

ROOT = Path(__file__).resolve().parents[1]
SLUG = "insolvenzforderungen-checker"
PLUGIN = ROOT / SLUG
CASE_SLUG = "insolvenzforderungen-handwerk-koeln"
CASE = ROOT / "testakten" / CASE_SLUG


def doc_text(path):
    with ZipFile(path) as archive:
        tree = ET.fromstring(archive.read("word/document.xml"))
    return " ".join(e.text or "" for e in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))


class Forderungspruefung(unittest.TestCase):
    def test_eleven_direct_skills_and_complete_routing(self):
        skills = {p.parent.name: p for p in (PLUGIN / "skills").glob("*/SKILL.md")}
        self.assertEqual(len(skills), 11)
        main = skills["forderungen-pruefen-und-tabelle-vorbereiten"].read_text()
        for name, path in skills.items():
            if name != "forderungen-pruefen-und-tabelle-vorbereiten":
                self.assertIn(f"`{name}`", main)
            text = path.read_text()
            for n in range(1, 7):
                self.assertRegex(text, rf"(?m)^## {n}\. ")
            self.assertIn("../../references/zitierweise.md", text)
            self.assertIn("Times New Roman 11 pt", text)
            self.assertNotRegex(name, r"werkstatt|schnellstart")

    def test_prompts_profiles_and_review_hashes(self):
        mini = (PLUGIN / f"{SLUG}-schnellstart.md").read_bytes()
        self.assertTrue(mini_within_limits(SLUG, mini))
        self.assertLessEqual(len(mini.decode()), 7500)
        self.assertEqual(validate_files(PLUGIN, SLUG, ROOT), [])
        review = json.loads((ROOT / "quality/evals" / f"{SLUG}.json").read_text())
        self.assertEqual({c["target_skill"] for c in review["cases"]}, {p.parent.name for p in (PLUGIN / "skills").glob("*/SKILL.md")})
        for key, suffix in (("mini_review", "schnellstart"), ("workshop_review", "werkstatt")):
            content = (PLUGIN / f"{SLUG}-{suffix}.md").read_bytes()
            self.assertEqual(review[key]["sha256"], hashlib.sha256(content).hexdigest())
            for anchor in ("IX ZR 114/23", "IX ZR 47/19", "IX ZR 195/01"):
                self.assertIn(anchor, content.decode())

    def test_electronic_handover_is_not_a_judicial_finding(self):
        content = (PLUGIN / "skills/tabelle-und-elektronische-uebergabe-vorbereiten/SKILL.md").read_text()
        for term in ("0300005", "044", "0300006", "3.6.2", "30.04.2027", "Freigabe"):
            self.assertIn(term, content)
        refs = (PLUGIN / "references/rechtsstand-und-entscheidungen.md").read_text()
        for term in ("sicheren Übermittlungsweg", "Paragraf 174 Absatz 4", "Paragraf 170", "Paragraf 175 SGB III", "Individualisierung", "Verteilungs"):
            self.assertIn(term, refs)
        helper_ref = (PLUGIN / "references/elektronische-uebergabe.md").read_text()
        self.assertIn("Schematron", helper_ref)
        self.assertIn("kein XJustiz", helper_ref)

    def test_native_inventory_and_no_answer_documents(self):
        files = [p for p in CASE.iterdir() if include_in_working_dump(p, CASE)]
        self.assertGreaterEqual(len(files), 55)
        self.assertTrue({".docx", ".eml", ".xlsx"} <= {p.suffix for p in files})
        self.assertEqual(len(list(CASE.glob("*_Anmeldung_*.docx"))), 10)
        self.assertFalse(any(p.suffix in (".md", ".py", ".json", ".yaml") for p in files))
        for path in files:
            self.assertTrue(path.name.isascii(), path.name)
            if path.suffix == ".docx":
                text = doc_text(path)
                self.assertGreater(len(text), 900, path.name)
                self.assertNotRegex(text, r"(?i)Musterlösung|Lösungsmatrix|Platzhalter|Formathinweis|Diese Testakte")
                self.assertNotIn("§", text)
                self.assertRegex(text, "[äöüÄÖÜß]", path.name)
        self.assertFalse((CASE / ".build_case.py").exists())

    def test_email_headers_and_csv_columns(self):
        for path in CASE.glob("*.eml"):
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            for key in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version"):
                self.assertTrue(message[key], f"{path.name}: {key}")
            sent = parsedate_to_datetime(message["Date"])
            self.assertLessEqual(sent.date(), date(2026, 9, 10))
            self.assertEqual(str(message["Date"]).split(",")[0], sent.strftime("%a"))
            self.assertGreater(len(message.get_body().get_content()), 500)
        for path in CASE.glob("*.csv"):
            rows = list(csv.reader(path.read_text(encoding="utf-8-sig").splitlines(), delimiter=";"))
            self.assertGreaterEqual(len(rows), 5)
            self.assertTrue(all(len(row) == len(rows[0]) for row in rows), path.name)

    def test_each_input_occurs_in_combined_pdf_without_warning(self):
        path = CASE / "gesamt-pdf" / f"{CASE_SLUG}_gesamt.pdf"
        pdf = PdfReader(path)
        self.assertGreater(len(pdf.pages), 55)
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        self.assertNotIn("Diese Testakte wurde", text)
        self.assertNotIn("This test case file", text)
        for p in CASE.iterdir():
            if include_in_working_dump(p, CASE):
                self.assertIn(p.name, text)

    def test_claim_amounts_and_chronology_from_original_records(self):
        amounts = ("6426", "17850", "11424", "4284", "2856", "14280", "48000", "23800", "60000", "8330")
        claims = sorted(CASE.glob("*_Anmeldung_*.docx"))
        for path, amount in zip(claims, amounts, strict=True):
            text = doc_text(path).replace(".", "").replace(" ", "")
            self.assertIn(amount, text, path.name)
        opening = doc_text(next(CASE.glob("12_Eroeffnungsbeschluss_*.docx")))
        for term in ("01.07.2026", "08:00", "14.08.2026", "15.09.2026"):
            self.assertIn(term, opening)
        account = doc_text(next(CASE.glob("15_Betriebskonto_*.docx")))
        credit = doc_text(next(CASE.glob("39_Gutschrift_*.docx")))
        self.assertIn("3570", account.replace(".", "").replace(" ", ""))
        self.assertIn("1190", credit.replace(".", "").replace(" ", ""))

    def test_workbook_totals_have_recalculated_values(self):
        path = next(CASE.glob("16_*.xlsx"))
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
        try:
            for sheet, cell, expected in (("Anmeldeeingang", "D17", 197250), ("Bankbelege", "D23", 142852)):
                self.assertTrue(formulas[sheet][cell].value.startswith("="))
                self.assertEqual(values[sheet][cell].value, expected)
        finally:
            formulas.close()
            values.close()


if __name__ == "__main__":
    unittest.main()
