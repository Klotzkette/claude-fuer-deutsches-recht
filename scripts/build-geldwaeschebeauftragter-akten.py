#!/usr/bin/env python3
"""Erzeugt die drei AML-Arbeitsakten aus dem kanonischen JSON, ohne Lösungen.

Excel, Gesamt-PDF und ZIPs entstehen in den getrennten Repository-Buildern.
Bestehende XLSX-Dateien und Gesamt-PDFs werden hier nicht verändert.
"""
from __future__ import annotations

import argparse
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime, formataddr
from html import escape
import json
import mimetypes
from pathlib import Path
import re
import zipfile
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'scripts/data/geldwaeschebeauftragter-akten.json'
NOTICE = ('Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\n'
          'This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.')


def font_paths():
    choices = [Path('/System/Library/Fonts/Supplemental'),
               Path('/usr/share/fonts/truetype/msttcorefonts'),
               Path('/usr/share/fonts/truetype/liberation2'),
               Path('/usr/share/fonts/truetype/liberation'), Path('/tmp/kk338/liberation')]
    for directory in choices:
        for regular, bold in [('Times New Roman.ttf', 'Times New Roman Bold.ttf'),
                              ('times.ttf', 'timesbd.ttf'),
                              ('LiberationSerif-Regular.ttf', 'LiberationSerif-Bold.ttf')]:
            found = list(directory.rglob(regular)) if directory.exists() else []
            if found and found[0].with_name(bold).exists():
                return found[0], found[0].with_name(bold)
    raise RuntimeError('Times New Roman oder Liberation Serif wird benötigt.')


def create_pdf(path, item):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    body = ParagraphStyle('Body', fontName='AkteSerif', fontSize=11, leading=14.5, spaceAfter=9)
    title = ParagraphStyle('Title', parent=body, fontName='AkteSerifBold', fontSize=16, leading=19, spaceAfter=16)
    small = ParagraphStyle('Small', parent=body, fontSize=9.5, leading=12, spaceAfter=0)
    def p(text, style=body):
        return Paragraph(escape(text).replace('\n', '<br/>'), style)
    story = [p(item['title'], title)]
    for text in item['paragraphs']:
        story.append(p(text))
    if item.get('table'):
        rows = [[p(str(v), small) for v in row] for row in item['table']]
        n = len(rows[0])
        widths = {2: [180, 307], 3: [150, 217, 120], 4: [81, 158, 145, 103]}[n]
        table = Table(rows, colWidths=widths, repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E9EDF1')),
            ('GRID', (0,0), (-1,-1), 0.4, colors.HexColor('#D9D9D9')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LEFTPADDING', (0,0), (-1,-1), 7), ('RIGHTPADDING', (0,0), (-1,-1), 7),
            ('TOPPADDING', (0,0), (-1,-1), 7), ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ]))
        story.extend([Spacer(1,5), table])
    def footer(canvas, document):
        canvas.setFont('AkteSerif',9)
        canvas.setFillColor(colors.HexColor('#606060'))
        canvas.drawString(54,29,item['id'])
        canvas.drawRightString(A4[0]-54,29,str(document.page))
    path.parent.mkdir(parents=True,exist_ok=True)
    SimpleDocTemplate(str(path),pagesize=A4,leftMargin=54,rightMargin=54,
                      topMargin=48,bottomMargin=48,title=item['title'],author='',
                      invariant=1).build(story,onFirstPage=footer,onLaterPages=footer)


def create_docx(path, item):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.oxml.ns import qn
    document=Document()
    section=document.sections[0]
    section.page_width=Mm(210); section.page_height=Mm(297)
    section.top_margin=Mm(20); section.bottom_margin=Mm(20)
    section.left_margin=Mm(22); section.right_margin=Mm(22)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=document.styles[name]
        style.font.name='Times New Roman'; style.font.size=Pt(11)
        style.font.color.rgb=RGBColor(0,0,0)
        rfonts=style._element.get_or_add_rPr().rFonts
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
            rfonts.attrib.pop(qn('w:'+attr),None)
        for attr in ['ascii','hAnsi','eastAsia','cs']:
            rfonts.set(qn('w:'+attr),'Times New Roman')
        for border in list(style._element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
        style.paragraph_format.space_after=Pt(9)
    document.styles['Title'].font.size=Pt(16)
    document.styles['Title'].font.bold=True
    document.styles['Title'].paragraph_format.space_after=Pt(16)
    document.styles['Heading 1'].font.bold=True
    document.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    document.styles['Heading 1'].paragraph_format.keep_with_next=True
    document.styles['Normal'].paragraph_format.line_spacing=1.08
    document.add_paragraph(item['title'],'Title')
    for text in item['paragraphs']:
        heading = bool(re.match(r'^\d+(?:\.\d+)*\s+\D.{0,80}$',text))
        paragraph=document.add_paragraph(text, 'Heading 1' if heading else 'Normal')
        paragraph.paragraph_format.widow_control=True
    document.core_properties.title=item['title']
    document.core_properties.author=''
    document.core_properties.last_modified_by=''
    document.core_properties.created=datetime(2026,10,8,8,0)
    document.core_properties.modified=datetime(2026,10,8,8,0)
    path.parent.mkdir(parents=True,exist_ok=True)
    document.save(path)
    # ZIP-Zeitstempel vereinheitlichen; die Native-Datei bleibt ein normales DOCX.
    with zipfile.ZipFile(path) as archive:
        entries=[(info.filename,archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
        for filename,content in entries:
            info=zipfile.ZipInfo(filename,(2026,10,8,8,0,0)); info.compress_type=zipfile.ZIP_DEFLATED
            archive.writestr(info,content)


def create_eml(case_dir, case, item):
    contacts={contact['id']:contact for contact in case['contacts']}
    docs={document['id']:document for document in case['documents']}
    def address(key):
        contact=contacts[key]
        return formataddr((contact['name'],contact['email']))
    message=EmailMessage(policy=SMTP)
    message['From']=address(item['sender'])
    message['To']=', '.join(address(key) for key in item['recipients'])
    stamp=datetime.fromisoformat(item['date']+'T'+item['time']).replace(tzinfo=ZoneInfo('Europe/Berlin'))
    message['Date']=format_datetime(stamp)
    message['Subject']=item['subject']
    message['Message-ID']=f"<{item['id'].lower()}-{item['date']}@aktenpost.example>"
    if item.get('reply_to'):
        previous=next(x for x in case['emails'] if x['id']==item['reply_to'])
        ref=f"<{previous['id'].lower()}-{previous['date']}@aktenpost.example>"
        message['In-Reply-To']=ref; message['References']=ref
    sender=contacts[item['sender']]
    signature=f"\n\n{sender['name']}\n{sender['role']}\n{sender['organisation']}\n{sender['address']}\n{sender['email']}"
    message.set_content(item['body']+signature,charset='utf-8')
    for identifier in item['attachments']:
        source=case_dir/docs[identifier]['path']
        mimetype=mimetypes.guess_type(source.name)[0] or 'application/octet-stream'
        maintype,subtype=mimetype.split('/',1)
        message.add_attachment(source.read_bytes(),maintype=maintype,subtype=subtype,filename=source.name)
    if item['attachments']:
        message.set_boundary('=_aml_'+item['id'].lower()+'_20261008')
    target=case_dir/item['path']; target.parent.mkdir(parents=True,exist_ok=True)
    target.write_bytes(message.as_bytes())


def readme(case_dir, case, release_tag):
    slug=case['slug']
    base='https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/'+release_tag
    text=f"# {case['title']}\n\n"
    text+=f"## 1. Auftrag und Ausgangslage\n\nBearbeitungsstand: **08.10.2026**. Ihre Rolle: {case['role']}. Die Akte enthält offene Tatsachen und widersprüchliche Angaben. Sie enthält keine Musterlösung und keine freigegebene externe Meldung.\n\n"
    text+='Die Bearbeitung soll den dokumentierten Sachstand, die Quellenherkunft und die jeweils noch benötigte Entscheidung auseinanderhalten. Erstellen Sie eine Belegmatrix, eine Übersicht der Beteiligten und Zahlungen, gezielte Rückfragen und ausformulierte interne Entscheidungsvorlagen.\n\n'
    text+='## 2. Bestand und Arbeit mit den Unterlagen\n\n'
    text+='Die Akte umfasst **25 Originaldateien**: acht PDF-Belege, vier bearbeitbare Word-Dokumente, zwölf echte E-Mail-Dateien mit binären Anlagen und eine Excel-Arbeitsmappe. Die E-Mail-Anlagen entsprechen den separat abgelegten Dateien. Rechnen Sie diese mehrfach vorliegenden Unterlagen nicht als zusätzliche Zahlungen oder zusätzliche Belege.\n\n'
    text+='Die Tabellen ordnen die Aktenangaben; eine geplante oder behauptete Zahlung bleibt von einem dokumentierten Geldfluss getrennt. Im Unternehmen und Notariat bestehen organisatorische Rückmeldetermine. Sie sind nicht als geprüfte gesetzliche Fristen vorgegeben.\n\n'
    text+='Alle Namen, Unternehmen, Personenbezüge, Anschriften, Kontokennungen und Vorgänge sind erfunden. Kontaktdaten verwenden ausschließlich reservierte `.example`-Domains; sie sind nicht für den Versand bestimmt. Register- und Bankunterlagen sind erkennbare Nachbildungen ohne echte Siegel oder verwendbare Zahlungsdaten.\n\n<!-- reserved-example-contacts -->\n\n'
    if slug=='aml-kanzlei-grundstueck-berlin':
        text+='Die Datei K-P08 stammt aus einer gesonderten vertraulichen Rechtsberatung. Ihr Inhalt ist gegenüber den übrigen Transaktionsunterlagen getrennt gekennzeichnet. Die Fiktion des politischen Amtes ist ausdrücklich dokumentiert; keine reale Amtsträgerin wird beschrieben.\n\n'
    text+='## 3. Dateiregister\n\n| Kennung | Datei | Dokumentherkunft |\n|---|---|---|\n'
    for item in case['documents']:
        text+=f"| {item['id']} | [{item['title']}]({item['path']}) | {item['origin']} |\n"
    for item in case['emails']:
        text+=f"| {item['id']} | [{item['subject']}]({item['path']}) | E-Mail vom {item['date']} |\n"
    text+='| XLSX | [Fallregister](04_Tabellen/Fallregister.xlsx) | Arbeitsmappe aus den Aktenangaben |\n\n'
    text+='## 4. Downloads\n\n'+NOTICE+'\n\n'
    text+='| Fassung | Download |\n|---|---|\n'
    text+=f"| Gesamt-PDF | [Akte am Stück](gesamt-pdf/{slug}_gesamt.pdf) |\n"
    text+=f"| Originaldateien | [Originalformat-ZIP]({base}/testakte-{slug}.zip) |\n"
    text+=f"| Einzel-PDFs | [Einzel-PDF-ZIP]({base}/testakte-{slug}-einzelpdfs.zip) |\n"
    (case_dir/'README.md').write_text(text,encoding='utf-8')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--release-tag',default='geldwaeschebeauftragter-v445.34.0')
    parser.add_argument('--case',action='append',dest='cases')
    parser.add_argument('--readme-only',action='store_true')
    args=parser.parse_args()
    data=json.loads(DATA.read_text(encoding='utf-8'))
    if not args.readme_only:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        regular,bold=font_paths()
        pdfmetrics.registerFont(TTFont('AkteSerif',str(regular)))
        pdfmetrics.registerFont(TTFont('AkteSerifBold',str(bold)))
        pdfmetrics.registerFontFamily('AkteSerif',normal='AkteSerif',bold='AkteSerifBold',italic='AkteSerif',boldItalic='AkteSerifBold')
    for case in data['cases']:
        if args.cases and case['slug'] not in args.cases:
            continue
        case_dir=ROOT/'testakten'/case['slug']; case_dir.mkdir(parents=True,exist_ok=True)
        if not args.readme_only:
            for item in case['documents']:
                path=case_dir/item['path']
                (create_pdf if path.suffix=='.pdf' else create_docx)(path,item)
            for item in case['emails']:
                create_eml(case_dir,case,item)
        readme(case_dir,case,args.release_tag)
        print(f"{case['slug']}: 8 PDF, 4 DOCX, 12 EML; Excel wird gesondert erzeugt.")


if __name__=='__main__':
    main()
