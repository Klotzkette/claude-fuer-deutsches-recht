#!/usr/bin/env python3
"""Zwei fiktive JVA-Arbeitsakten aus kanonischen Inhalten erzeugen.

DOCX/PDF verwenden Times New Roman 11 pt. XLSX werden über die gebündelte
Artifact-Tool-Laufzeit mit Formeln erstellt, nachgerechnet und gerendert.
Vorhandene JVA-Akten werden nicht gelesen oder überschrieben.
"""
from pathlib import Path
import argparse, datetime as dt, hashlib, json, os, re, subprocess, sys
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
QA = Path('/tmp/jva-erweiterungsakten-qa')
RUNTIME = Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies'
NOTICE = 'Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\nThis test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'

def paras(s): return [p.strip() for p in s.strip().split('\n\n') if p.strip()]
def doc(c, n, ext, title, text):
    c['docs'].append(dict(file=f'{n:02d}_{title.replace(" ", "_")}.{ext}', title=title.replace('_',' '), paragraphs=paras(text)))

def hessen():
    c=dict(slug='jva-hessen-absonderung-hagedorn',person='Jannis Hagedorn',title='Jannis Hagedorn Absonderung und Aufklärung in Hessen',prison='JVA Waldhain',record='WH-418/26',land='Hessen',lawyer='Rechtsanwältin Malin Rehbein',lawmail='m.rehbein@kanzlei-rehbein.example',domain='jva-waldhain.example',docs=[],mails=[])
    doc(c,1,'docx','Mandatsauftrag_16_Oktober', '''Kanzlei Rehbein, Fiktiver Auensteg 14, 34121 Kassel. Aktenzeichen RH 214/26. Vermerk vom 16.10.2026, 09.20 Uhr.

1. Auftrag

Jannis Hagedorn, geboren am 19.04.1984, verbüßt seit dem 03.02.2025 eine rechtskräftige Freiheitsstrafe in der erfundenen JVA Waldhain, Fiktive Waldhainstraße 80, 34125 Kassel. Der Vorgang betrifft Erwachsenenstrafvollzug in Hessen. Es liegt weder Untersuchungshaft noch Maßregelvollzug vor. Frau Rehbein übernimmt heute die Vertretung vom Kollegen Volker Brack. Die Vollmacht liegt vor. Die Aktenbezeichnung der Strafvollstreckungskammer lautet ausdrücklich fiktiv: ÜBUNG 7 StVK 214/26.

Herr Hagedorn will die seit dem 09.09.2026 bestehende Trennung von den übrigen Gefangenen überprüfen lassen. Er berichtet von Schmerzen nach dem Zugriff am 09.09.2026 und von immer kürzeren Gesprächen an der Tür. Er verlangt nicht, ohne jede Prüfung auf die bisherige Station zurückzukehren. Einen stufenweisen Kontaktplan mit anderer Gruppe und begleiteten Gesprächen würde er versuchen. Seine Schwester möchte verhindern, dass der Streit mit einzelnen Bediensteten jede andere Frage überlagert.

2. Bisheriger Verfahrensstand

Der frühere Rechtsanwalt hat am 22.09.2026 einen Antrag mit Hauptsachebegehren und vorläufigem Begehren übermittelt. Die Kammer hat am 24.09.2026 ausschließlich den vorläufigen Antrag abgelehnt. Die Hauptsache ist nach Auskunft der Geschäftsstelle noch anhängig. Das Gerichtspostfach weist für den Beschluss eine elektronische Zustellung am 25.09.2026 aus. Die Kopie an den Gefangenen wurde erst am 28.09.2026 ausgegeben. Das sind verschiedene Ereignisse; eine Fristberechnung ist bislang nicht dokumentiert.

3. Erwartetes Arbeitsergebnis

Die Kanzlei soll alle Maßnahmen, Zeitabschnitte, Angaben und Belege auseinanderhalten und den vorhandenen Entwurf vervollständigen. Vor einer weiteren Einreichung sind der Stand der Hauptsache, die Art des bisherigen gerichtlichen Beschlusses, die Bekanntgabe der Fortdauerentscheidung und die vollständigen ärztlichen Unterlagen zu klären. Die Befürchtung eines Übergriffs ist eine unbewiesene Darstellung. Die Angaben der Bediensteten sollen genauso konkret ausgewertet werden. Ein Ermittlungsauftrag oder eine Strafanzeige ist noch nicht erteilt.

Die beiden Tabellen sind Arbeitsunterlagen. Zeitdifferenzen ersetzen weder die Prüfung des zutreffenden Landesrechts noch die Berechnung einer Rechtsbehelfsfrist. Es soll kein Entwurf versandt werden, bevor Frau Rehbein den Text, Anlagen, aktuellen medizinischen Sachstand und Empfänger persönlich freigegeben hat.''')
    doc(c,2,'docx','Vollmacht_und_Kontaktwunsch', '''Jannis Hagedorn, WH-418/26, JVA Waldhain. Erklärung vom 15.10.2026.

Ich bevollmächtige Rechtsanwältin Malin Rehbein, mich wegen der Absonderung, der damit verbundenen Unterbringungsbedingungen und der gerichtlichen Überprüfung dieser Maßnahmen gegenüber der JVA und den zuständigen Gerichten zu vertreten. Sie darf die hierzu erforderlichen Akten einsehen, Auskünfte anfordern und Erklärungen entgegennehmen. Eine Zustimmung zu einem Vergleich oder zur Rücknahme eines Antrags möchte ich erst nach einem persönlichen Gespräch erteilen.

Ich bin damit einverstanden, dass die Anstalt die für diesen Vorgang erforderlichen medizinischen Befunde unmittelbar an Frau Rehbein übermittelt. Meine Schwester Marlene darf organisatorische Nachrichten austauschen. Sie soll nicht ohne weitere Rücksprache meine vollständige Gesundheitsakte bekommen. Über Untersuchungen und Behandlungen möchte ich weiterhin selbst mit dem medizinischen Dienst sprechen.

Die Post an meine Schwester darf nicht mit anwaltlicher Korrespondenz verwechselt werden. Ich benutze in der Anstalt keinen frei zugänglichen E-Mail-Anschluss. Wenn in dieser Akte eine E-Mail meiner Schwester steht, stammt sie von ihr. Meine eigenen Erklärungen werden als Brief oder als unterschriebener Gesprächsvermerk übermittelt.

Die bisherige Kanzlei Brack darf die vorhandenen Verfahrensunterlagen an Frau Rehbein herausgeben. Die Gründe für den Wechsel sollen sachlich bleiben. Ich bin mit Herrn Brack unzufrieden, weil ich den eingereichten Text vorher nicht gelesen hatte. Ich kann nicht beurteilen, welche seiner rechtlichen Aussagen zutreffen.

Jannis Hagedorn, eigenhändig unterzeichnet am 15.10.2026. Die Unterschrift wird in dieser fiktiven Datei nur durch den ausgeschriebenen Namen wiedergegeben.''')
    doc(c,3,'docx','Eingabe_an_die_Anstaltsleitung', '''Jannis Hagedorn, WH-418/26, an die Leitung der JVA Waldhain. 18.09.2026.

Sehr geehrte Frau Dr. Gründel,

ich beantrage, die Trennung von den anderen Gefangenen zu beenden. Falls Sie eine Rückkehr auf Station 3 derzeit ausschließen, bitte ich um eine andere Gruppe oder einen zeitlich gestuften Versuch. Ich möchte wieder mit jemandem zusammen essen und nicht nur durch die geöffnete Tür sprechen. Bitte teilen Sie mir schriftlich mit, was ich konkret tun muss, damit die Einschränkung überprüft und gelockert wird.

Am 09.09.2026 war ich wütend. Ich hatte den Brief über den Krankenhausaufenthalt meines Vaters erst am Nachmittag erhalten. Ich habe laut gesprochen und wollte nicht zurück in den Haftraum. Dass ich einen Stuhl gegen jemanden geworfen hätte, stimmt nach meiner Erinnerung nicht. Ich habe ihn zur Seite geschoben. Ich räume ein, dass der andere Gefangene zurückgewichen ist und dass meine Worte bedrohlich geklungen haben können.

Beim anschließenden Zugriff wurde mein linker Arm festgehalten. Seitdem tut die Schulter weh. Ich habe keine Fotos. Ich bitte um eine Untersuchung und darum, den Verlauf zu dokumentieren. Mir wurde gesagt, der Arzt habe mich gesehen. Ich erinnere mich an eine Frage an der Tür, aber nicht an eine Untersuchung meiner Schulter. Vielleicht wird hier ein Gespräch mit etwas anderem verwechselt.

Am 12.09.2026 wollte ich nicht allein in den Hof. Am 13. und 14.09.2026 habe ich darum gebeten und nach meiner Erinnerung keine Gelegenheit bekommen. Ich möchte, dass diese Tage einzeln geprüft werden. Eine einmalige Ablehnung bedeutet nicht, dass ich grundsätzlich auf Bewegung verzichte. Ich möchte auch meiner Schwester erklären können, wie es weitergeht.

Mit freundlichen Grüßen
Jannis Hagedorn''')
    doc(c,4,'docx','Gespraech_in_der_JVA_15_Oktober', '''Kanzlei Rehbein. Gespräch am 15.10.2026 von 14.05 bis 14.48 Uhr im anwaltlichen Besuchsraum der JVA Waldhain. Anwesend: Jannis Hagedorn und Malin Rehbein. Niederschrift am selben Abend.

Herr Hagedorn sitzt ruhig und beantwortet die Fragen zusammenhängend. Frau Rehbein stellt keine medizinische Diagnose. Er berichtet, dass der Vater mittlerweile wieder zu Hause sei. Die Schulterbeschwerden seien schwächer, beim Anziehen merke er sie noch. Er habe einer Pflegekraft zweimal davon erzählt. An Namen könne er sich nur bei Frau Eschenbach erinnern. Die Anwältin hat keine sichtbaren Verletzungen fotografiert und keinen körperlichen Befund erhoben.

Zur Situation am 09.09.2026 erklärt er: Er habe gesagt, man könne ihn nicht immer herumkommandieren. Ob er dabei mit dem Stuhl aufgestanden sei, könne er nach dem Streit nicht sicher trennen. Er bestreitet einen gezielten Angriff. Die Anwältin hält diese Unsicherheit fest und bittet ihn, Erinnerungen nicht zu ergänzen, nur weil andere Berichte etwas Bestimmtes behaupten.

Herr Hagedorn nennt drei längere Kontakte: ein Gespräch mit dem Sozialdienst am 17.09., ein Telefonat mit Marlene am 29.09. und das heutige anwaltliche Gespräch. Kurze Kontakte bei Essenausgabe und Reinigung habe es täglich gegeben. Er schildert sie als hastig und wechselnd. An einem Tag habe ihm Herr Bering länger zugehört. Daraus lasse sich aber keine verlässliche tägliche Gesprächszeit rekonstruieren.

Die Anwältin erläutert, dass eine bestehende Gefährdung nicht durch die bloße Bezeichnung als Strafe widerlegt wird. Für den nächsten Text werden deshalb die ursprüngliche Anordnung, ihre Fortdauer, die tatsächlichen Kontaktmöglichkeiten und die gesundheitlichen Folgen getrennt behandelt. Herr Hagedorn akzeptiert einen begleiteten Versuch außerhalb der bisherigen Station. Er möchte nicht, dass sein früherer Mitgefangener ungefragt mit seinen Gesundheitsangaben konfrontiert wird.

Die Notiz wurde Herrn Hagedorn vorgelesen. Er berichtigt, das Telefonat sei am 29.09. und nicht am 28.09. gewesen. Weitere Korrekturen hat er nicht genannt. Eine unterschriebene Rückbestätigung steht noch aus.''')
    doc(c,5,'docx','Schreiben_Kanzlei_Brack_21_September', '''Rechtsanwalt Volker Brack, Fiktive Steinwiese 6, 34119 Kassel, an die JVA Waldhain. 21.09.2026. Zeichen VB 88/26.

Sehr geehrte Damen und Herren,

Herr Hagedorn berichtet mir, dass er seit einiger Zeit allein gehalten wird. Ich fordere Sie auf, die Einzelhaft sofort aufzuheben. Mein Mandant ist nicht damit einverstanden. Nach dem Strafvollzugsgesetz des Bundes ist eine solche Behandlung nur ausnahmsweise zulässig. Die Anstalt wird sich nicht auf ihre Hausordnung berufen können, wenn mein Mandant seine Rechte geltend macht.

Ich habe die Eingabe meines Mandanten vom 18.09.2026 gelesen. Die dort geschilderten Schmerzen sind ernst zu nehmen. Ob eine Untersuchung stattgefunden hat, müssen Sie wissen. Ich gehe davon aus, dass sich der Sachverhalt in Ihren Akten findet. Eine gesonderte Aufstellung der einzelnen Tage halte ich deshalb derzeit nicht für nötig. Bitte sehen Sie von weiterem Schriftverkehr über Formalien ab und schicken Sie meinen Mandanten zurück auf die bisherige Station.

Sollte ich bis morgen Mittag keine positive Nachricht haben, werde ich das Gericht einschalten. Mein Mandant kann sich auf die bestehenden Regeln zum Strafvollzug berufen. Ich werde auch die Kosten der Sache geltend machen. Bitte bestätigen Sie mir, dass alle Maßnahmen aufgehoben werden.

Mit freundlichen Grüßen
Volker Brack, Rechtsanwalt''')
    doc(c,6,'docx','Gerichtsantrag_22_September', '''Rechtsanwalt Volker Brack. An die Strafvollstreckungskammer. 22.09.2026. Fiktives Aktenzeichen nach Eingangsmitteilung: ÜBUNG 7 StVK 214/26.

1. Antrag

Ich beantrage für Herrn Jannis Hagedorn, die Einzelhaft aufzuheben und ihn sofort auf seine bisherige Station zurückzubringen. Dies wird sowohl im Hauptverfahren als auch im Eilverfahren verlangt. Die Behandlung ist unmenschlich und durch nichts gerechtfertigt. Die Kosten müssen der Anstalt auferlegt werden.

2. Begründung

Mein Mandant wird schon länger allein gehalten. Er hatte einen Streit, aber das rechtfertigt nicht, ihn auf unbestimmte Zeit wegzuschließen. Die Behörden müssen das Strafvollzugsgesetz des Bundes beachten. Die Gefängnisleitung weiß, dass mein Mandant Beschwerden hat. Er hat auch einen Arzt verlangt. Die genauen Unterlagen hat die Anstalt und nicht der Gefangene.

Die Eingabe meines Mandanten wird als Anlage beigefügt. Ich mache ihren Inhalt zum Gegenstand meines Vortrags. Eine genaue Liste der einzelnen Anordnungen liegt mir nicht vor. Es kann keinen Unterschied machen, ob die Anstalt von Absonderung oder Einzelhaft spricht. Die Kammer wird die Sache nach Aktenlage entscheiden können. Wegen der fortdauernden Situation muss sofort entschieden werden.

3. Anlagen

Beigefügt ist die Eingabe vom 18.09.2026. Die Anordnung vom 09.09.2026 wurde angefordert, lag bei Versand dieses Schriftsatzes jedoch noch nicht vor. Weitere Belege werden nachgereicht, soweit dies erforderlich erscheint.

Volker Brack, Rechtsanwalt. Versandprotokoll der Kanzlei: 22.09.2026, 16.41 Uhr; elektronische Eingangsbestätigung desselben Tages, 16.42 Uhr.''')
    doc(c,7,'docx','Uebergabe_der_Vorakte', '''Kanzlei Brack an Kanzlei Rehbein. Übergabevermerk vom 16.10.2026, 08.10 Uhr.

Die Vertretung wird entsprechend dem Wunsch des Mandanten beendet. Ich übergebe die Eingabe vom 18.09., mein Schreiben vom 21.09., den Gerichtsantrag vom 22.09., die Entscheidung vom 24.09. und die nachträglich eingegangenen Anstaltsunterlagen. Die ursprüngliche Anordnung lag bei Fertigung des Gerichtsantrags nicht vor. Sie kam am 23.09. über die Schwester. Ich habe keine weitere Begründung an das Gericht gesandt.

Die Fortdauerentscheidung vom 09.10. wurde mir am 12.10. als Kopie übersandt. Wann der Mandant sie vollständig erhielt, habe ich nicht mit ihm klären können. In meiner Papierakte steht am Rand „Beschwerde?“. Eine Rechtsbeschwerde ist nicht eingereicht worden. Eine gesonderte Prüfung, ob gegen den vorläufigen Beschluss überhaupt ein solcher Rechtsbehelf gegeben ist, liegt in der Akte nicht vor.

Die ärztliche Stellungnahme vom 12.09. wurde nur in der knappen Weitergabe des medizinischen Dienstes angekündigt. Einen vollständigen Befund habe ich nicht. Ich habe die Vorgänge nicht an die Staatsanwaltschaft gegeben. Der Mandant hatte zunächst vor allem die Rückkehr auf seine Station verlangt. Ein gesonderter Auftrag wegen eines möglichen körperlichen Übergriffs war nicht eindeutig erteilt.

Die Kostenrechnung über 238,00 EUR wurde bezahlt. Eine nachträgliche Gebührenerweiterung ist nicht vorgesehen. Frau Marlene Hagedorn hat ausschließlich organisatorisch geholfen. Sie war nicht meine Auftraggeberin. Bitte bestätigen Sie den Eingang und teilen Sie mit, wenn in der Dateiübertragung ein Anhang fehlt.

Volker Brack''')
    doc(c,8,'docx','Fortfuehrung_des_Verfahrens_Entwurf', '''Rechtsanwältin Malin Rehbein. An die zuständige Strafvollstreckungskammer des Landgerichts. Fiktives Verfahrenszeichen ÜBUNG 7 StVK 214/26. Entwurf vom 16.10.2026, nicht eingereicht.

1. Verfahrensgegenstand und Anträge

Ich zeige die Vertretung von Herrn Jannis Hagedorn an. Die Vollmacht ist beigefügt. Der Hauptsacheantrag vom 22.09.2026 wird nach Abgleich mit dem vollständigen Verwaltungsvorgang wie folgt präzisiert: [Die konkret angegriffene Anordnung, ihre bekannt gegebenen Fortsetzungen und die jeweils beantragte Rechtsfolge nach abschließender Akteneinsicht einsetzen.] Mein Mandant verlangt eine individuelle Entscheidung über die fortdauernde Absonderung und über praktisch geeignete Schritte aus dieser Situation. Eine Rückkehr ohne jede Sicherheitsprüfung ist nicht das einzige von ihm akzeptierte Ergebnis.

Der vorläufige Antrag wird unter Berücksichtigung des Beschlusses vom 24.09.2026 und der seitdem entstandenen Umstände gesondert begründet: [Das vorläufige konkrete Begehren, dessen tatsächliche Dringlichkeit und den derzeitigen medizinischen Stand einsetzen.] Dieser Text bezeichnet keine unstatthafte Rechtsbeschwerde gegen einen Beschluss nach § 114 Absatz 2 StVollzG. Vor Einreichung ist festzuhalten, ob eine neue vorläufige Regelung, eine Änderung aufgrund nachträglicher Umstände oder ein anderer verfahrensrechtlicher Schritt beantragt wird.

2. Chronologie und belegte Veränderungen

Die Anstalt ordnete am 09.09.2026 eine Absonderung an. Ihr Bekanntgabevermerk nennt 18.15 Uhr. Die Eingabe vom 18.09.2026 enthält sowohl das Begehren nach gemeinschaftlicher Unterbringung als auch die Bitte um Untersuchung der Schulter. Die frühe gerichtliche Eingabe hat diese Gegenstände nicht getrennt dargestellt. Der Beschluss vom 24.09.2026 betrifft nach seinem Tenor nur das vorläufige Begehren. Belege: A1 bis A5.

Am 09.10.2026 wurde die Fortdauer angeordnet. Die Anstalt nennt eine weiterhin mögliche Auseinandersetzung mit zwei bestimmten Mitgefangenen. Die Notiz des Sozialdienstes erwähnt dagegen einen begleiteten Kontakt in anderer Zusammensetzung. Es ist nicht geklärt, ob eine solche Maßnahme konkret geprüft, verworfen oder bereits versucht wurde. Ebenso fehlt der vollständige Nachweis der aufsichtsbehördlichen Befassung. Belege: A6 und A7. [Bekanntgabe, angeforderte vollständige Genehmigung und Stellungnahme der Abteilung ergänzen.]

3. Rechtlicher Ausgangspunkt und offene Subsumtion

Für den Erwachsenenstrafvollzug ist das hessische Landesrecht anhand seiner am Maßnahmetag geltenden Fassung anzuwenden. Besondere Sicherungsmaßnahmen und das Verfahren sind an §§ 50 und 51 HStVollzG zu prüfen. Aus dem bloßen Etikett „Absonderung“ folgt weder die Rechtmäßigkeit noch die Rechtswidrigkeit. Der anfängliche Anlass, die späteren Tatsachen, die Betreuung, ärztliche Beteiligung, Dokumentation und die Entscheidung über Fortdauer müssen jeweils am konkreten Zeitabschnitt geprüft werden. [Den verifizierten amtlichen Fassungsstand und die für die einzelnen Zeitabschnitte einschlägigen Absätze ergänzen.]

Der verfassungsrechtliche Maßstab ersetzt diese Tatbestandsprüfung nicht. Die Entscheidung des Bundesverfassungsgerichts vom 20.04.2026, Az. 2 BvR 1650/24, Rn. 32–46, betrifft den Maßregelvollzug und eine nicht hinreichend begründete Verfassungsbeschwerde. Sie darf hier nicht als unmittelbare Entscheidung über hessischen Strafvollzug ausgegeben werden. Ihre Erwägungen zur fortlaufenden Rechtfertigung länger dauernder Absonderung sind anhand der unterschiedlichen Regelungsgrundlagen und des konkreten Verlaufs zu prüfen. [Den tragenden Vergleich und seine Grenzen anhand des amtlichen Volltextes ausformulieren.]

Der Hinweis auf fehlende Unterlagen ist kein Beweis für einen Übergriff. Andererseits kann eine nicht beigefügte ärztliche Stellungnahme die streitige Frage einer Untersuchung nicht ohne weiteres beantworten. Es wird um Beiziehung des vollständigen Sicherungsvorgangs, der konkret bezeichneten ärztlichen Unterlagen und der maßgeblichen Dienstberichte gebeten. Gesundheitsdaten anderer Personen sind nicht Gegenstand dieses Antrags. Belege: A8 bis A10.

4. Gegenargumente und Aufklärung

Das dokumentierte aggressive Verhalten und die Sorge vor erneuten Auseinandersetzungen werden ernst genommen. Die Anstalt soll erläutern, weshalb auch ein begleiteter Kontakt mit anderen Personen derzeit ausscheidet und woran die Fortdauerentscheidung überprüfbar anknüpft. Die behauptete Erkrankung des Vaters erklärt die Erregung, entschuldigt aber keinen Angriff. Umgekehrt ist ein unter Druck entstandener Streit kein dauerhafter Ersatz für die weitere Einzelfallprüfung.

Vor einer abschließenden Fassung fehlen die vollständige medizinische Stellungnahme, das Ergebnis der Schulteruntersuchung, die Antwort zur Aufsichtsbehörde und ein belastbarer aktueller Kontaktplan. [Die heute fortdauernden Belastungen, das beantragte vorläufige Mindestmaß an Kontakt und den Umgang mit widersprechenden Berichten konkret ergänzen.] Der Antragsteller bietet an, seine Erinnerungen in einem geordneten Gespräch zu präzisieren. Eine bereits feststehende Misshandlung wird nicht behauptet.

5. Anlagen und Freigabe

A1 Anordnung 09.09.; A2 Bekanntgabe; A3 Eingabe 18.09.; A4 Antrag 22.09.; A5 Beschluss 24.09.; A6 Fortdauer 09.10.; A7 Nachricht Sozialdienst; A8 Frühdienstbericht; A9 Spätdienstbericht; A10 Vermerk des medizinischen Dienstes; A11 Gespräch 15.10.; A12 Vollmacht. Die Anlagen sind vor Einreichung auf Vollständigkeit und lesbare Seitennummern zu prüfen. Die übrigen Unterlagen verbleiben zunächst in der Arbeitsakte.

Malin Rehbein, Rechtsanwältin. Bearbeitungsfassung ohne Versandfreigabe.''')
    doc(c,9,'pdf','Anordnung_09_September', '''JVA Waldhain, Fiktive Waldhainstraße 80, 34125 Kassel. WH-418/26. Entscheidung vom 09.09.2026, 17.30 Uhr, an Jannis Hagedorn.

1. Entscheidung

Herr Hagedorn wird ab heute von den übrigen Gefangenen abgesondert und im Haftraum S 4 untergebracht. Die Entscheidung wird als besondere Sicherungsmaßnahme getroffen. Sie ist keine Disziplinarentscheidung. Die Abteilungsleitung überprüft die Fortdauer nach Eingang der beteiligten Dienstberichte. Aufenthalt im Freien wird getrennt organisiert. Anwaltliche Besuche und notwendige medizinische Versorgung bleiben möglich.

2. Gründe

Am Nachmittag geriet Herr Hagedorn mit zwei Mitgefangenen in einen lautstarken Streit. Nach der ersten Meldung trat er gegen einen Stuhl, nahm ihn anschließend hoch und bewegte sich auf einen Mitgefangenen zu. Der hinzugerufene Dienst beendete die Situation. Ob der Stuhl tatsächlich zum Schlag eingesetzt werden sollte, steht nach den bislang vorliegenden Angaben nicht fest. Zwei Beteiligte erklärten, sie würden ihm derzeit nicht gemeinsam begegnen wollen.

Die Abteilungsleitung befürchtet eine erneute Auseinandersetzung bei unmittelbarer Rückkehr. Bis zur Klärung sind Kontakte außerhalb der bisherigen Gruppe zu prüfen. Ein automatischer Ausschluss sämtlicher Gespräche wird nicht angeordnet. Der Sozialdienst wird um einen ersten Gesprächstermin gebeten. Beschwerden über Schmerzen sind an den medizinischen Dienst weiterzugeben.

Die Maßnahme wird auf die Regelungen zu besonderen Sicherungsmaßnahmen im Hessischen Strafvollzugsgesetz gestützt. Über Dauer, Betreuung und ärztliche Beteiligung ist gesondert fortlaufend zu entscheiden. Dr. Ottilie Gründel, Leiterin der Vollzugsabteilung.''')
    doc(c,10,'pdf','Bekanntgabe_und_Postlauf', '''JVA Waldhain, Stationsdienst, Vermerk WH-418/26 vom 10.09.2026.

Die schriftliche Anordnung vom 09.09.2026 wurde Herrn Hagedorn am selben Tag um 18.15 Uhr vollständig ausgehändigt und erklärt. Herr Hagedorn wollte den Empfang nicht unterschreiben. Die Aushändigung wurde durch Levin Bering und Hiltrud Bär bestätigt. Aus der verweigerten Unterschrift wurde kein Einverständnis oder Rechtsmittelverzicht abgeleitet.

Herr Hagedorn bat am 10.09. um eine weitere Kopie für seine Schwester. Diese ging in die reguläre Ausgangspost. Eine direkte elektronische Zustellung an eine Kanzlei erfolgte zu diesem Zeitpunkt nicht. Die Kanzlei Brack meldete sich erst später. Der Vermerk enthält keinen Nachweis eines Gerichtseingangs.

Nachtrag vom 13.10.2026: Für die Fortdauerentscheidung vom 09.10. liegt auf der Stationsliste nur der Eintrag „Kopie abgegeben, Mittag“ ohne Empfangsunterschrift vor. Der Dienst kann den genauen Zeitpunkt derzeit nicht bestätigen. Ein weiterer Ausgabevermerk wurde angefordert. Die Kanzlei erhielt am 12.10. eine PDF-Kopie per E-Mail. Diese Kopie beantwortet nicht die Frage, wann der Gefangene die vollständige Ausfertigung erhalten hat.

Hiltrud Bär, Vollzugsdienst.''')
    doc(c,11,'pdf','Dienstbericht_Fruehdienst', '''JVA Waldhain. Ergänzender Bericht von Levin Bering vom 10.09.2026 zu seinem Einsatz am 09.09., 14.10 bis 14.25 Uhr.

Ich wurde über den Stationsruf wegen eines lautstarken Streits hinzugezogen. Beim Betreten des Gemeinschaftsraums stand Herr Hagedorn seitlich am Tisch. Der Stuhl vor ihm war gekippt. Ich sah nicht, wie er in diese Lage gekommen war. Herr Hagedorn hielt zu diesem Zeitpunkt keinen Gegenstand in der Hand. Der Kollege Ewald Stieg erklärte, der Stuhl sei zuvor angehoben worden.

Ich forderte Herrn Hagedorn auf, mit mir aus dem Raum zu gehen. Er sprach laut und ging einen Schritt zurück. Anschließend wurde er von zwei Kollegen an den Armen geführt. Ich sicherte den Weg und hatte auf die genaue Griffhaltung nicht durchgehend freie Sicht. Ein gezielter Schlag eines Bediensteten wurde von mir nicht gesehen. Damit treffe ich keine Aussage über den gesamten Zeitraum vor meinem Eintreffen.

Im Vorraum sagte Herr Hagedorn, seine Schulter tue weh. Ich gab diese Information an die Dienstleitung weiter. Der medizinische Dienst wurde nach meiner Erinnerung telefonisch verständigt. Ob die Person anschließend untersucht wurde, weiß ich nicht. Die Notiz in der Übergabeliste „Arzt informiert“ stammt von mir und bedeutet keine Bestätigung einer abgeschlossenen Untersuchung.

Die Bezeichnung als Frühdienst bezieht sich auf meine Schicht von 06.00 bis 14.30 Uhr. Der Bericht wurde nach Dienstende begonnen und am nächsten Morgen abgeschlossen. Levin Bering.''')
    doc(c,12,'pdf','Dienstbericht_Spaetdienst', '''JVA Waldhain. Bericht von Ewald Stieg vom 09.09.2026, 20.15 Uhr. Spätdienst Station 3.

Herr Hagedorn rief einem Mitgefangenen zu, er solle sich nicht einmischen. Er hatte einen Stuhl angehoben und bewegte ihn vor seinem Körper. Ich stand am anderen Ende des Tisches. Ob er den Stuhl gegen die Person richten oder Platz schaffen wollte, konnte ich in diesem Moment nicht sicher unterscheiden. Ich ging von einer akuten Gefahr aus und forderte Unterstützung an.

Nachdem der Stuhl wieder auf dem Boden stand, wurden beide Arme des Gefangenen gehalten. Herr Hagedorn zog den linken Arm zurück und drehte sich zur Seite. Wir führten ihn aus dem Raum. Die Haltung wurde im Vorraum gelöst. In der ersten Kurzmeldung habe ich „Angriffsversuch mit Stuhl“ geschrieben. Eine absichtliche Körperverletzung ist damit nicht gerichtlich festgestellt.

Ich habe den Eintrag „medizinisch abgeklärt“ aus der Übergabe übernommen. Eine eigene Untersuchung oder ein ärztliches Gespräch habe ich nicht beobachtet. Der Gefangene äußerte nach der Verbringung, er brauche einen Arzt. Ich informierte Frau Bär. Den späteren Ablauf kenne ich nicht.

Auf Station 3 waren zwei Mitgefangene sehr erregt. Ein gemeinsamer Aufenthalt noch am selben Abend erschien mir nicht möglich. Zur Frage einer anderen Gruppe an späteren Tagen kann ich aus diesem Einsatzbericht keine Aussage treffen. Ewald Stieg.''')
    doc(c,13,'pdf','Mitteilung_medizinischer_Dienst', '''Medizinischer Dienst der JVA Waldhain. Auskunft an die Vollzugsabteilung vom 08.10.2026, WH-418/26. Verfasst von Pflegefachperson Mira Eschenbach.

In der Übernahmeliste vom 09.09. steht „Schulter links, Arzt informieren“. Für den 10.09. ist ein kurzer Kontakt am Haftraum dokumentiert. Die handschriftliche Notiz ist nicht vollständig lesbar. Eine strukturierte Untersuchung der Schulter lässt sich aus dieser Liste allein nicht erkennen. Die Anforderung einer ärztlichen Stellungnahme zur Sicherungsmaßnahme ist von einer Befunddokumentation über eine Verletzung zu unterscheiden.

Eine ärztliche Stellungnahme zur weiteren Unterbringung soll am 12.09. erstellt worden sein. Sie liegt in der medizinischen Teilakte. In der an die Vollzugsabteilung versandten Kopie fehlen die zweite Seite und die Unterschrift. Die medizinische Geschäftsstelle wurde um Vervollständigung gebeten. Es lässt sich aus dieser unvollständigen Kopie nicht ableiten, welche konkreten Untersuchungen erfolgt sind.

Herr Hagedorn sprach am 30.09. erneut über Beschwerden bei bestimmten Bewegungen. Ein Termin wurde vorgemerkt. Ob dieser Termin durchgeführt wurde, muss anhand des Behandlungsjournals geprüft werden. Akute Veränderungen sollen unabhängig von der rechtlichen Auseinandersetzung sofort medizinisch bewertet werden. Dieser Vermerk gibt keine Diagnose und keine Empfehlung zur Beendigung oder Fortsetzung der Sicherungsmaßnahme ab.

An die bevollmächtigte Kanzlei dürfen die angeforderten personenbezogenen Unterlagen nach Prüfung der vorgelegten Einwilligung unmittelbar übermittelt werden. Eine Übermittlung an Angehörige ist damit nicht verbunden.''')
    doc(c,14,'pdf','Brief_Marlene_Hagedorn', '''Marlene Hagedorn, Fiktiver Fasanenrain 22, 34225 Baunatal. An meinen Bruder Jannis, JVA Waldhain. 04.10.2026.

Lieber Jannis,

Papa ist seit Donnerstag wieder zu Hause. Er kann noch nicht lange telefonieren, aber er hat sich über Deine Karte gefreut. Ich habe ihm nicht jedes Detail aus Deinen Briefen erzählt. Er fragt ohnehin schon, ob er Dir mehr helfen müsste. Bitte sag mir beim nächsten Gespräch, ob ich ihm Deinen letzten Brief vorlesen darf.

Beim Telefonat am 29.09. warst Du so kurz angebunden, dass ich erst dachte, Du seist auf mich böse. Du hast gesagt, Du wüsstest nicht mehr, welcher Tag was gewesen sei. Ich kann Deine Schulter nicht beurteilen und will auch niemandem etwas unterstellen, was ich nicht gesehen habe. Aber ich möchte, dass es vernünftig untersucht und aufgeschrieben wird. Ich habe die Kanzlei um eine Erklärung gebeten, was sie jetzt eigentlich beantragt hat.

Den neuen Pullover konnte ich nicht einfach abgeben. Die Pforte hat mir das Verfahren erklärt; ich war wohl mit der falschen Erwartung hingefahren. Den Kassenbon hebe ich auf. Es ist nichts Dringendes, Du sollst nur wissen, warum das Paket noch nicht da ist.

Wenn ein Gespräch in einer anderen Gruppe hilft, probiere es bitte. Du musst dabei niemandem recht geben. Ich möchte nur, dass die Wochen nicht alle gleich aussehen. Ich hoffe, wir bekommen bald einen Besuchstermin. Marlene.''')
    doc(c,15,'pdf','Fortdauerentscheidung_09_Oktober', '''JVA Waldhain, WH-418/26. Entscheidung vom 09.10.2026. An Jannis Hagedorn.

1. Entscheidung

Die am 09.09.2026 angeordnete Absonderung wird fortgesetzt. Eine erneute Überprüfung durch die Abteilungsleitung ist für den 19.10.2026 vorgesehen. Die Maßnahme ist vorher zu ändern, wenn ihre Voraussetzungen entfallen. Die Organisation eines begleiteten Einzelgesprächs mit dem Sozialdienst wird aufrechterhalten.

2. Begründung

Zwei frühere Mitgefangene haben erklärt, dass eine gemeinsame Rückkehr derzeit Spannungen auslösen würde. Herr Hagedorn hat im Gespräch am 06.10. erneut erklärt, er lasse sich nicht vor anderen zurechtweisen. Die Abteilung bewertet dies als fortbestehendes Risiko. Ein erneuter körperlicher Vorfall während der Absonderung ist nicht dokumentiert. Die wiederholte Ablehnung einzelner Hofzeiten wird berücksichtigt; der Gefangene bestreitet, generell darauf verzichtet zu haben.

Der Sozialdienst hat einen Kontaktversuch außerhalb der bisherigen Gruppe vorgeschlagen. Über Raum, Begleitung und Teilnehmer liegt noch keine abschließende Abstimmung vor. Die Anstalt sieht deshalb bislang keine unmittelbar verfügbare Alternative. Die Zustimmung der Aufsichtsbehörde wurde nach dem internen Vermerk am 08.10. angefragt. Eine Kopie der Antwort ist dem vorliegenden Bescheid nicht beigefügt. Die Anstaltsleitung hat eine weitere Sachstandsabfrage veranlasst.

Die medizinische Stellungnahme wird im Verwaltungsvorgang als vorhanden geführt. Der Vollständigkeit der übermittelten Kopie ist nachzugehen. Eine medizinisch gebotene sofortige Anpassung wird durch diese Entscheidung nicht ausgeschlossen. Dr. Ottilie Gründel.''')
    doc(c,16,'pdf','Beschluss_24_September', '''Fiktives Gerichtsdokument. Strafvollstreckungskammer des Landgerichts. Aktenzeichen ausschließlich für die Übungsakte: ÜBUNG 7 StVK 214/26. Beschluss vom 24.09.2026.

1. Tenor

Der Antrag auf vorläufige Beendigung der Absonderung und sofortige Rückverlegung auf Station 3 wird abgelehnt. Die Entscheidung über den Hauptsacheantrag und die Kosten des Hauptverfahrens bleibt vorbehalten.

2. Gründe

Nach dem schriftlichen Antrag ist es zu einem Streit mit weiteren Gefangenen gekommen. Die Anstalt hat telefonisch mitgeteilt, dass eine Eskalation verhindert werden musste und die Trennung fortlaufend überprüft werde. Ein schriftlicher Bericht ist von der Kammer nicht angefordert worden. Die Eingabe nennt weder eine bestimmte ärztliche Diagnose noch eine aktuelle konkrete Gefahr aus der Unterbringung.

Die begehrte Rückkehr würde das Ergebnis des Hauptverfahrens vorwegnehmen. Die Kammer sieht für eine solche Anordnung nach dem derzeitigen Vortrag keine ausreichende Grundlage. Der Gefangene kann sich an den medizinischen Dienst wenden. Der gerichtliche Antrag lässt nicht erkennen, weshalb die organisatorische Abstimmung der Anstalt nicht abgewartet werden kann. Eine andere Form der vorläufigen Kontaktgestaltung ist nicht ausdrücklich beantragt.

Die Kammer hat ohne vorherige schriftliche Anhörung der Anstalt und ohne Beiziehung der medizinischen Teilakte entschieden. Eine Sachentscheidung über die längerfristige Fortdauer wird damit nicht getroffen. Der Hauptsachevorgang wird nach Vorlage der Verwaltungsakte weiterbearbeitet.

Die Unterschrift wird in dieser fiktiven Akte durch „Dr. Thekla Morgenstern, Richterin“ wiedergegeben. Es handelt sich nicht um einen Beschluss eines wirklichen Gerichts und nicht um eine zitierbare Rechtsprechungsquelle.''')
    doc(c,17,'pdf','Gerichtspost_Eingangsnachweise', '''Kanzlei Brack. Ausdruck des elektronischen Postjournals, erstellt am 16.10.2026. Bezug ÜBUNG 7 StVK 214/26.

Der Gerichtsantrag wurde am 22.09.2026 um 16.41 Uhr an das gerichtliche Postfach übermittelt. Der technische Eingang wurde um 16.42 Uhr bestätigt. Das Journal verzeichnet ein Hauptdokument und eine Anlage. Der Anhang ist die Eingabe des Gefangenen vom 18.09.2026. Die spätere Anordnung der JVA war nicht Bestandteil dieses Sendevorgangs.

Der Beschluss vom 24.09.2026 ging am 25.09.2026 um 09.12 Uhr ein. Das elektronische Empfangsbekenntnis wurde von Rechtsanwalt Brack mit Zustelldatum 25.09.2026 abgegeben. Die Weiterleitung an Herrn Hagedorn wurde am selben Tag als Brief vorbereitet. Die Stationsliste weist eine Aushändigung dieser Kopie am 28.09.2026 aus.

Für die Hauptsache ist weder ein abschließender Beschluss noch eine Rücknahme verzeichnet. Ein ausgehender Rechtsbehelf gegen den Beschluss vom 24.09.2026 ist nicht gespeichert. Dieser Ausdruck belegt die genannten Vorgänge im Kanzleisystem, berechnet aber keine Frist und erklärt keinen Rechtsbehelf für statthaft.

Die technische Aufstellung ist vollständig fiktiv. Es wurden keine echten Postfachkennungen, Zertifikate oder Signaturdaten verwendet.''')
    doc(c,18,'pdf','Hausordnung_und_Kontaktangebot', '''JVA Waldhain, interner Arbeitsauszug der Abteilung S vom 01.08.2026. Ausgegeben am 23.09.2026. Kein amtlicher Gesetzes- oder Verordnungstext.

Bei getrennten Hofzeiten teilt der diensthabende Vollzugsdienst der betroffenen Person das verfügbare Zeitfenster persönlich mit. Eine Ablehnung wird mit Datum, Uhrzeit und dem mitgeteilten Grund dokumentiert. Ein später geäußerter Wunsch wird an die Dienstleitung gegeben. Aus der Ablehnung eines einzelnen Angebots folgt kein allgemeiner Verzicht auf weitere Angebote.

Der Sozialdienst wird bei länger dauernder Absonderung über die tatsächlichen Kontaktmöglichkeiten unterrichtet. Kurze Versorgungskontakte werden von vereinbarten Gesprächen getrennt verzeichnet. Die Besuchsabteilung stimmt familiäre Besuche mit der zuständigen Abteilung ab. Eine Anmeldung durch Angehörige ist noch keine Zusage.

Wenn eine Person Beschwerden nach einer Sicherungsmaßnahme berichtet, wird der medizinische Dienst informiert. Die Information ersetzt keine Befunddokumentation. Anordnungen, Fortdauerentscheidungen und die tatsächliche Durchführung sind in den jeweiligen Vorgängen zu dokumentieren. Diese interne Organisationsbeschreibung erweitert keine gesetzliche Eingriffsbefugnis.

Kontakt: Vollzugsabteilung, vollzug@jva-waldhain.example. Die Adresse ist für den fiktiven Vorgang reserviert.''')
    doc(c,19,'pdf','Rechnung_und_Zahlungsnotiz', '''Kanzlei Brack, Fiktive Steinwiese 6, 34119 Kassel. Rechnung VB-2026-104 vom 02.10.2026 an Jannis Hagedorn.

Berechnet wird die vereinbarte außergerichtliche Erstbearbeitung einschließlich Aktenanlage und schriftlicher Erstberatung. Pauschalhonorar netto 200,00 EUR, Umsatzsteuer 19 Prozent 38,00 EUR, Gesamtbetrag 238,00 EUR. Die Rechnung enthält keine Abrechnung eines weiteren gerichtlichen Rechtsbehelfs. Grundlage ist die gesonderte Honorarvereinbarung vom 20.09.2026, die im Original in der Kanzlei verbleibt und auf Anforderung übergeben werden soll.

Zahlungsnotiz vom 07.10.2026: 238,00 EUR eingegangen, Verwendungszweck Hagedorn VB 88/26. Die Zahlung wurde von Marlene Hagedorn für ihren Bruder ausgelegt. Daraus wurde kein eigenes Mandat der Schwester abgeleitet. Eine Rückzahlung oder Kostenfestsetzung ist bislang nicht beantragt.

Die Kanzlei verwendet in dieser Übungsakte keine echte Kontoverbindung. Für eine Kostenprüfung fehlen der vollständige Auftrag und die Honorarvereinbarung. Die Rechnung ist ein Sachverhaltsbeleg und keine Aussage darüber, ob die Vergütung im Einzelfall wirksam vereinbart oder erstattungsfähig ist.''')
    doc(c,20,'pdf','Raumblatt_und_Uebergaben', '''JVA Waldhain, Raumblatt S 4, Stand 01.09.2026. Interne technische Unterlage.

Der Haftraum ist nach dem Bestandsplan 3,80 m lang und 2,70 m breit. Die darin enthaltene Sanitärzone ist mit 1,55 m² angegeben. Der Plan benennt eine Gesamtraumfläche von 10,26 m². Es gibt ein Bett, einen Tisch, einen Stuhl und einen Schrank. Der Sanitärbereich ist durch eine raumhohe Wand und eine Tür vom übrigen Raum getrennt. Ein aktualisiertes Aufmaß nach der Leitungsreparatur fehlt.

Die Belegung ist im fraglichen Zeitraum als Einzelbelegung dokumentiert. Dieses Raumblatt belegt keine tägliche Einschlussdauer, keine Luftqualität und keine tatsächliche Nutzung sämtlicher Gegenstände. Angaben zu Hofgang, Gesprächen und Bewegungszeiten werden im Stationsjournal geführt.

Übergabe vom 08.10.2026, 14.00 Uhr: Bitte prüfen, ob die Liste der Kontakte auch Telefonate enthält. Der bisherige Ausdruck zählt nur Angebote des Sozialdienstes. Die Abteilungsleitung benötigt den tatsächlichen Verlauf für die nächste Entscheidung. Ein Antrag des Gefangenen auf Einsicht in die vollständige Tagesdokumentation liegt vor.

Gero Tann, technischer Dienst, und Hiltrud Bär, Stationsdienst. Der Text ist keine Bescheinigung einer rechtlich ausreichenden Mindestfläche.''')
    return c

def bw():
    c=dict(slug='jva-bw-schutzmassnahmen-finkbeiner',person='Amira Finkbeiner',title='Amira Finkbeiner Schutz und Unterbringung in Baden Württemberg',prison='JVA Silberrain',record='SR-227/26',land='Baden-Württemberg',lawyer='Rechtsanwalt Xaver Lindenfels',lawmail='x.lindenfels@kanzlei-lindenfels.example',domain='jva-silberrain.example',docs=[],mails=[])
    doc(c,1,'docx','Mandatsauftrag_16_Oktober', '''Kanzlei Lindenfels, Fiktive Rosenhalde 9, 72074 Tübingen. Aktenzeichen XL 307/26. Übernahmevermerk vom 16.10.2026, 08.45 Uhr.

1. Auftrag und Person

Amira Finkbeiner, geboren am 08.11.1993, befindet sich seit dem 02.06.2026 im Erwachsenenstrafvollzug der erfundenen JVA Silberrain, Fiktiver Silberrain 40, 72076 Tübingen. Eine Untersuchungshaft oder Unterbringung nach dem Psychisch-Kranken-Hilfe-Recht liegt nicht vor. Ihr Ehemann Idris Finkbeiner lebt mit ihrer Tochter Nora in Rottenburg. Die Mandantin spricht Deutsch und Arabisch; für die derzeitigen Gespräche wünscht sie keine Übersetzung. Diese Entscheidung kann sie jederzeit ändern.

Am 28.09.2026 entstand eine ernst zu nehmende psychische Krise. Die Akte beschreibt weder eine Selbstschädigungsmethode noch eine Anleitung. Die Mandantin will, dass ihr damaliger Zustand nicht verharmlost wird. Sie beanstandet die Dauer einzelner Einschränkungen, den Umgang mit ihren Nachfragen und die anschließende enge Mehrfachunterbringung. Medizinische Behandlung soll fortgeführt werden. Die anwaltliche Arbeit ersetzt weder ärztliche Beurteilung noch sofortige Hilfe bei einer akuten Gefährdung.

2. Verfahrensstand

Die bisherige Vertreterin Sigrun Kappel reichte am 02.10.2026 einen als Beschwerde bezeichneten Schriftsatz bei der Strafvollstreckungskammer ein. Die Kammer behandelte das Begehren als Antrag auf vorläufigen Rechtsschutz und als Antrag auf gerichtliche Entscheidung. Am 09.10.2026 wurde nur der vorläufige Antrag abgelehnt. Die Zustellung an die Kanzlei ist für den 12.10.2026 nachgewiesen. Am 13.10. wurde eine sofortige Beschwerde an dieselbe Kammer gesandt. Über deren Behandlung ist noch nichts bekannt. Das fiktive gerichtliche Zeichen lautet ÜBUNG 4 StVK 307/26.

3. Bearbeitung

Zu prüfen sind die einzelnen Maßnahmen und Zeitabschnitte: Unterbringung im besonders gesicherten Haftraum vom 28.09. bis 01.10., Entzug bestimmter Gegenstände, Versorgung mit geeignetem Ersatzmaterial, Beobachtung, ärztliche Beteiligung sowie die spätere Doppel- und vorübergehende Dreifachbelegung. Der tatsächliche Verlauf ist wichtiger als eine zusammenfassende Bezeichnung als Einzelhaft. Die Raumunterlagen nennen unterschiedliche Flächen; Originale und Bedeutungen sind zu erhalten.

Erwartet werden eine nachvollziehbare Ereignis- und Belegübersicht, die Prüfung der bisherigen Prozessführung und ein ausformulierter, nach Rückfragen zu vervollständigender Schriftsatz. Aktuelle Schutzbedürftigkeit, Würde, Privatheit und die Belastung anderer Gefangener sind gemeinsam zu betrachten. Vor einer Einreichung entscheidet Herr Lindenfels über Text, Rechtsweg, Anlagen und Freigabe.''')
    doc(c,2,'docx','Vollmacht_und_Vertraulichkeit', '''Amira Finkbeiner, SR-227/26, JVA Silberrain. Erklärung vom 15.10.2026.

Ich bevollmächtige Rechtsanwalt Xaver Lindenfels, mich wegen der Schutzmaßnahmen seit dem 28.09.2026, der anschließenden Unterbringung und der hierzu geführten gerichtlichen Verfahren zu vertreten. Er darf die erforderlichen Akten und Befunde anfordern und mich gegenüber der Anstalt und den Gerichten vertreten. Eine Rücknahme, ein Vergleich oder eine Erklärung, ich sei mit der gesamten Behandlung einverstanden gewesen, bedarf vorher meiner Zustimmung.

Ich willige ein, dass der medizinische Dienst die für diesen Vorgang erforderlichen Befunde und Anordnungen an meinen Anwalt übermittelt. Meine Tochter soll keine Einzelheiten über die Krise bekommen. Mein Mann darf Termine abstimmen und meine Briefe weitergeben. Über die Weitergabe medizinischer Einzelheiten an ihn möchte ich jeweils mitentscheiden. Die Zustimmung zu seiner Hilfe ist keine pauschale Freigabe meiner Gesundheitsakte.

Ich habe in der Anstalt keinen allgemeinen Zugang zu E-Mail. Die in der Akte enthaltenen elektronischen Nachrichten werden von der Kanzlei, der Anstalt oder meinen Angehörigen versandt. Meine eigenen Erklärungen sind Briefe oder Gesprächsnotizen. Ich möchte vor einer gerichtlichen Einreichung den fertigen Text vorgelesen bekommen und Fehler berichtigen können.

Die bisherige Kanzlei Kappel darf die vollständige Vorgangsakte einschließlich Sende- und Empfangsnachweisen übergeben. Ich bin mit meiner Behandlung nicht in allen Punkten einverstanden, bestreite aber nicht, dass ich Ende September dringend Unterstützung brauchte.

Amira Finkbeiner. Eigenhändige Unterzeichnung am 15.10.2026; in dieser fiktiven Datei nur als ausgeschriebener Name wiedergegeben.''')
    doc(c,3,'docx','Eingabe_30_September', '''Amira Finkbeiner, SR-227/26, an die Leitung der JVA Silberrain. 30.09.2026.

Sehr geehrte Frau Seibold,

ich bitte darum, dass heute noch mit mir besprochen wird, welche Schutzmaßnahmen weiterhin nötig sind und was sich ändern kann. Ich weiß, dass Sie sich am Montag ernsthaft Sorgen machen mussten. Ich möchte weiter mit dem ärztlichen und psychologischen Dienst sprechen. Ich bitte nicht darum, mich ohne Hilfe allein zu lassen.

Der Gürtel und die normalen Bettlaken wurden mir abgenommen. Ich habe besondere Kleidung und eine andere Decke bekommen. In der ersten Nacht war mir kalt. Am nächsten Morgen wurde eine zusätzliche Decke ausgegeben. Bitte prüfen Sie, ob die Versorgung von Anfang an genügt hat und welche Bettwäsche ich nun benutzen darf. Ich möchte nicht immer nur die Antwort bekommen, man entscheide das später.

Ich möchte auch verstehen, wer mich beobachtet und welche Bereiche des Raums einsehbar sind. Die Notwendigkeit von Kontrolle kann ich nachvollziehen. Trotzdem schäme ich mich bei der Körperpflege. Bitte erklären Sie mir, wie meine Privatsphäre soweit möglich geschützt wird. Ich möchte diese Frage in Ruhe mit einer Mitarbeiterin besprechen.

Wenn ich den jetzigen Raum verlassen kann, bitte ich um einen ruhigen Platz. Ein gemeinsamer Raum könnte helfen, wenn es mit der anderen Frau passt. Ich möchte aber weder sie belasten noch die ganze Nacht durch Gespräche wach bleiben. Bitte fragen Sie sie ebenfalls. Mit freundlichen Grüßen, Amira Finkbeiner.''')
    doc(c,4,'docx','Anwaltliches_Gespraech_15_Oktober', '''Kanzlei Lindenfels. Gespräch am 15.10.2026 von 10.10 bis 11.02 Uhr mit Amira Finkbeiner im Besuchsraum der JVA Silberrain. Vermerk von Xaver Lindenfels.

Frau Finkbeiner berichtet, dass sie inzwischen wieder regelmäßige Gespräche mit dem psychologischen Dienst führt. Sie möchte die Behandlung fortsetzen. Zu ihrem Zustand am 28.09. sagt sie, sie sei damals verzweifelt gewesen und habe nicht mehr richtig zuhören können. Sie hält es für nachvollziehbar, dass die Mitarbeiter sofort reagiert haben. Über einzelne Gefährdungsdetails wird in diesem anwaltlichen Vermerk bewusst nicht geschrieben.

Die Mandantin erinnert die erste Nacht als sehr kalt. Das Stationsprotokoll nennt die Ausgabe einer Schutzdecke um 19.30 Uhr und einer zweiten Decke am nächsten Morgen. Sie kann nicht sicher sagen, ob es zwischendurch eine weitere Ausgabe gab. Sie möchte, dass dies anhand der Materialliste geprüft wird. Sie sagt nicht, die erste Decke sei überhaupt nicht vorhanden gewesen.

Seit dem 01.10. lebt sie im Raum B 12 mit einer weiteren Frau. Vom Abend des 03.10. bis zum Morgen des 06.10. war eine dritte Person dort. Die Mandantin beschreibt Streit um Ruhezeiten, eine kaum abschließbare Sanitärtür und ein zugestelltes Tischende. Sie hatte einer Unterbringung mit einer Person zugestimmt, nachdem man ihr eine ruhigere Mitgefangene angekündigt hatte. Eine Zustimmung zur dritten Person erinnert sie nicht. Der Anwalt hat die beiden anderen Frauen nicht befragt.

Die Raumgröße kennt Frau Finkbeiner nicht. Sie hat ihrem Mann einmal „vielleicht zwölf Quadratmeter“ gesagt. Dies sei geschätzt gewesen. Die Zahl aus der früheren Kanzlei habe sie erst im Schriftsatz gesehen. Sie will keine falschen Maße behaupten. Besonders belastend seien die langen Abende im Raum und die Unterbrechungen der Nachtruhe gewesen.

Frau Finkbeiner berichtet außerdem, sie habe sich bei einer Begleitung am 29.09. grob behandelt gefühlt. Sie kann weder einen Namen noch eine genaue Uhrzeit nennen. Sie möchte vorerst eine sachliche Dokumentation und ein vertrauliches Gespräch. Ein konkreter körperlicher Übergriff wird in diesem Vermerk nicht festgestellt. Die Frage ist getrennt von der medizinischen Notwendigkeit der Schutzmaßnahme zu klären.''')
    doc(c,5,'docx','Kanzleischreiben_01_Oktober', '''Rechtsanwältin Sigrun Kappel, Fiktiver Mühlrain 13, 72070 Tübingen, an die JVA Silberrain. 01.10.2026. Zeichen SK 155/26.

Sehr geehrte Damen und Herren,

ich vertrete Frau Finkbeiner. Mein Mandat betrifft die Unterbringung und die Abnahme persönlicher Gegenstände. Ich verlange, dass meine Mandantin umgehend in eine normale Einzelzelle verlegt wird und ihre Sachen zurückerhält. Der Ehemann hat mir erklärt, dass es ihr wieder besser gehe. Es besteht deshalb kein Grund, sie weiterhin in dieser Weise zu behandeln.

Nach meiner bisherigen Kenntnis stehen Gefangenen mindestens sieben Quadratmeter allein zur Verfügung. Die Unterbringung mit zwei anderen Frauen ist deshalb von vornherein unzulässig. Ich werde mich nicht darauf einlassen, dass Sie bei der Fläche erst die Sanitärzone abziehen oder wieder dazurechnen. Die Behörde muss eine ordentliche Zelle bereitstellen. Der Mann meiner Mandantin schätzt den Raum auf zwölf Quadratmeter.

Soweit Sie sich auf Suizidgefahr berufen, bitte ich um Vorlage einer aktuellen schriftlichen Begründung. Ein Hinweis auf den Zustand vom Montag kann nicht alle späteren Maßnahmen erklären. Ich kann die Gesundheitsakte derzeit nicht vorlegen. Die Mandantin hat mir telefonisch noch nicht alle Einzelheiten schildern können. Das soll die gebotene Abhilfe nicht verzögern.

Bitte bestätigen Sie die Umsetzung bis morgen. Andernfalls werde ich Beschwerde einlegen. Mit freundlichen Grüßen, Sigrun Kappel, Rechtsanwältin.

Kanzleinotiz vom 02.10.: Die Information über eine dritte Person war eine Ankündigung aus einem Telefonat. Ihr tatsächlicher Einzug wurde bei Fertigung dieses Schreibens noch nicht bestätigt.''')
    doc(c,6,'docx','Beschwerdeschrift_02_Oktober', '''Rechtsanwältin Sigrun Kappel. An die Strafvollstreckungskammer. 02.10.2026. Nachträglich zugeordnetes fiktives Zeichen ÜBUNG 4 StVK 307/26.

1. Beschwerde und Begehren

Für Frau Amira Finkbeiner lege ich Beschwerde gegen die Behandlung in der JVA Silberrain ein. Ich beantrage, der Anstalt unverzüglich aufzugeben, alle besonderen Maßnahmen aufzuheben, sämtliche Gegenstände herauszugeben und eine Einzelzelle mit mindestens sieben Quadratmetern zur Verfügung zu stellen. Die Beschwerde soll sofort und auch in der Hauptsache entschieden werden.

2. Gründe

Die Mandantin befand sich am Montag in einer schwierigen Lage. Mittlerweile ist nach Auskunft ihres Ehemanns eine Beruhigung eingetreten. Die Fortdauer lässt sich deshalb nicht rechtfertigen. Die Wegnahme von Bettwäsche und Gürtel sowie die Beobachtung greifen schwer in ihre Rechte ein. Die Anstalt kann nicht einfach behaupten, sie müsse die Mandantin schützen, ohne den Schutzbedarf zu belegen.

Weiter ist eine gemeinschaftliche Unterbringung in einem ungefähr zwölf Quadratmeter großen Raum geplant oder bereits erfolgt. Eine weitere Person soll hinzukommen. Nach der von mir zugrunde gelegten Sieben-Quadratmeter-Grenze ist das unzulässig. Die genaue Aufteilung des Raums ist deshalb nicht entscheidend. Die Mandantin hat einer solchen Enge nicht zugestimmt.

Eine vollständige Chronologie ist noch nicht vorhanden. Die Anstalt möge alle Unterlagen vorlegen. Ich gehe davon aus, dass das Gericht die sofortige Beschwerde nach den allgemeinen Vorschriften entscheidet. Die Rechte meiner Mandantin müssen ohne weitere Verzögerung durchgesetzt werden.

3. Anlagen

Beigefügt sind die Vollmacht aus der Vorakte und das Schreiben an die Anstalt vom 01.10.2026. Der Raumplan, medizinische Vermerke und Einzelanordnungen liegen der Kanzlei noch nicht vor. Sigrun Kappel, Rechtsanwältin.''')
    doc(c,7,'docx','Aktenuebergabe_Kappel', '''Kanzlei Kappel an Kanzlei Lindenfels. Vermerk vom 16.10.2026, 08.00 Uhr. Akte SK 155/26.

Die Mandantin hat den Wechsel bestätigt. Ich übergebe die Eingaben vom 01. und 02.10., die gerichtliche Entscheidung vom 09.10. und die ergänzenden Unterlagen. Den Raumplan erhielt ich erst am 07.10. Meine Flächenangabe im ersten Schreiben beruhte auf einer telefonischen Schätzung des Ehemanns und einer älteren Notiz aus einer anderen Sache. Eine Entscheidung mit der dort genannten starren Grenze habe ich in dieser Akte nicht abgelegt.

Das Gericht hat den Schriftsatz nach seiner Mitteilung als Hauptsacheantrag und Antrag auf vorläufige Regelung behandelt. Der vorläufige Antrag wurde abgelehnt. Ich habe am 13.10.2026 eine sofortige Beschwerde an die Kammer geschickt. Eine vertiefte Prüfung der Statthaftigkeit ist nicht dokumentiert. Eine Entscheidung des Oberlandesgerichts liegt nicht vor. Die Hauptsache ist nach meinem Kenntnisstand weiter offen.

Die vollständige ärztliche Dokumentation wurde angefordert. Die Einwilligung war zunächst nur für organisatorische Auskünfte formuliert und wurde am 14.10. präzisiert. Es ist deshalb möglich, dass eine verzögerte Übersendung auf diese Rückfrage zurückgeht. Ich habe keine medizinische Aussage dazu getroffen, ob die Mandantin jetzt ohne jede weitere Beobachtung sicher wäre.

Die Mandantin hat mich gebeten, die Frage einer als grob empfundenen Begleitung zunächst zurückzustellen. Ich habe keine Strafanzeige veranlasst. Die Unterlagen enthalten die Namen der beiden Mitgefangenen nur gekürzt. Eine ungefragte Weitergabe ihrer medizinischen Daten soll vermieden werden.

Sigrun Kappel. Die Rechnung über 357,00 EUR ist nach Kanzleibuchung noch nicht ausgeglichen.''')
    doc(c,8,'docx','Ergaenzung_und_Rechtsschutz_Entwurf', '''Rechtsanwalt Xaver Lindenfels. An die zuständige Strafvollstreckungskammer. Fiktives Zeichen ÜBUNG 4 StVK 307/26. Arbeitsentwurf vom 16.10.2026, nicht eingereicht.

1. Verfahrensklärung

Ich zeige die Übernahme der Vertretung von Amira Finkbeiner an. Die Hauptsacheeingabe vom 02.10.2026 und die nachfolgende sofortige Beschwerde sind nach ihrem tatsächlichen Rechtsschutzziel und ihrer Statthaftigkeit zu prüfen. [Den aktuellen gerichtlichen Verfahrensstand, den Umgang mit der Eingabe vom 13.10. und das konkrete heute zu verfolgende Begehren nach Akteneinsicht einsetzen.] Ein Beschluss nach § 114 Absatz 2 StVollzG wird durch seine Bezeichnung als sofortige Beschwerde nicht regulär anfechtbar.

Die Mandantin beanstandet mehrere selbständige Maßnahmen und Zeitabschnitte. Sie bestreitet den Schutzbedarf am 28.09.2026 nicht. Sie verlangt eine nachvollziehbare Prüfung der Art, Dauer und tatsächlichen Ausgestaltung jeder Einschränkung. Eine Entscheidung über medizinisch gebotene Hilfe soll durch den vorliegenden Antrag nicht verhindert werden. [Den passenden Hauptsacheantrag für noch fortdauernde und bereits erledigte Maßnahmen jeweils ausformulieren; ein Feststellungsinteresse nicht nur behaupten.]

2. Ablauf und Unterbringung

Die besonders gesicherte Unterbringung dauerte vom 28.09.2026, 18.40 Uhr, bis zum 01.10.2026, 10.20 Uhr. Diese Zeit ist vom späteren Aufenthalt im Raum B 12 zu trennen. Der Gürtel blieb nach der Folgeverfügung zunächst im Verwahrbereich. Schutzmaterial wurde ausgegeben; streitig ist, ob in der ersten Nacht ausreichendes Material rechtzeitig verfügbar war. Die spätere Ausgabe gewöhnlicher Bettwäsche erfolgte laut Liste am 04.10. Die genaue medizinische Begründung und die individuellen Überprüfungen sind vollständig beizuziehen. Belege: A1 bis A4.

Vom 01. bis 03.10. war B 12 mit zwei Personen belegt. Eine dritte Person kam am Abend des 03.10. hinzu und verließ den Raum am Morgen des 06.10. Der Bestandsplan nennt 14,20 Quadratmeter einschließlich Sanitärzone. Das Aufmaß nennt 13,86 Quadratmeter bis zur veränderten Innenwand. Beide Zahlen sind mit ihrer jeweiligen Messgrundlage zu prüfen; eine Schätzung des Ehemanns ist kein Aufmaß. Belege: A5 und A6.

3. Landesrecht und gerichtlicher Maßstab

Die materielle Prüfung richtet sich nach dem für Erwachsenenstrafvollzug einschlägigen dritten Buch des baden-württembergischen Justizvollzugsgesetzbuchs. § 67 unterscheidet die einzelnen besonderen Sicherungsmaßnahmen; § 68 betrifft die dort geregelte Einzelhaft. Die bloße Unterbringung in einem Einzelraum ist damit nicht gleichzusetzen. Für Anordnung und ärztliche Beteiligung sind §§ 70 und 71 gesondert heranzuziehen. Für die spätere gemeinsame Unterbringung ist § 13 zu prüfen. Die Anstalt soll die tatsächliche Ausnahmegrundlage, die Dauer und ihre konkrete Entscheidung erläutern.

Das Bundesverfassungsgericht verlangt in seinem Beschluss vom 22.03.2016, Az. 2 BvR 566/15, Rn. 17–24, eine konkrete Befassung mit den Umständen der Unterbringung. Der Entscheidung wird hier keine bundesweit starre Quadratmetergrenze entnommen. Fläche, Sanitärabtrennung, Belegungsdauer, Einschlusszeiten und Ausweichmöglichkeiten sind deshalb mit ihren Belegen vorzutragen. [Aktuelle Raummaße, Beschaffenheit der Sanitärtür und tatsächliche tägliche Öffnungszeiten nach Aufklärung einsetzen.]

Für den vorläufigen Rechtsschutz sind Dringlichkeit und erreichbare Regelung konkret zu benennen. Der Beschluss des Bundesverfassungsgerichts vom 30.09.2024, Az. 2 BvR 150/24, Rn. 29–38, ist mit seinem tatsächlichen Verfahrenskontext auszuwerten. Er ersetzt keine Prüfung, welcher Schritt nach dem hier schon ergangenen vorläufigen Beschluss noch statthaft und sinnvoll ist. [Neue Umstände seit dem 09.10., aktueller medizinischer Stand und angemessenes vorläufiges Begehren ergänzen.]

4. Schutzbedarf und weniger belastende Gestaltung

Die ärztliche Notiz beschreibt eine ernsthafte Krise und eine zunächst notwendige enge Begleitung. Diese Angaben dürfen nicht allein mit der privaten Einschätzung des Ehemanns widerlegt werden. Die Mandantin wünscht fortgesetzte Behandlung und Gespräche. Zu prüfen bleibt, welche Gegenstände wann unter welchen Bedingungen wieder ausgegeben werden konnten, ob geeignetes Ersatzmaterial ausreichend vorhanden war und wie Schutz und Privatheit miteinander verbunden wurden.

Die vorübergehende Aufnahme einer dritten Frau wird mit einem Wasserschaden begründet. Das kann organisatorischen Handlungsbedarf erklären, beantwortet aber nicht die tatsächliche Nutzbarkeit des Raums und die individuelle Zumutbarkeit. Der Vortrag soll weder eine Mitgefangene für die Krise verantwortlich machen noch deren Gesundheitsdaten offenlegen. [Konkrete Alternative, verfügbare Räume und verbleibende Schutzbelange der Anstalt ergänzen.]

5. Anlagen und Abschluss

A1 Schutzanordnung; A2 Bekanntgabe; A3 Beobachtungsvermerk; A4 ärztliche Mitteilung; A5 Raumblatt; A6 Belegungsmitteilung; A7 Eingabe 30.09.; A8 frühere gerichtliche Eingabe; A9 Beschluss 09.10.; A10 elektronischer Zustellnachweis; A11 Gespräch 15.10.; A12 Vollmacht. Materiallisten und vollständige Behandlungsunterlagen sind angefordert und noch nicht vollständig vorhanden. Kein Anlagenverweis darf eine tatsächlich fehlende Datei als vorhanden darstellen.

Xaver Lindenfels, Rechtsanwalt. Der Text enthält ausdrücklich offene Tatsachen- und Antragsfelder. Die Mandantin muss die fertige Tatsachendarstellung vor Einreichung prüfen können.''')
    doc(c,9,'pdf','Schutzanordnung_28_September', '''JVA Silberrain, Fiktiver Silberrain 40, 72076 Tübingen. SR-227/26. An Amira Finkbeiner. 28.09.2026, 18.40 Uhr.

1. Anordnung

Die Gefangene wird vorläufig in einem besonders gesicherten Haftraum untergebracht. Gürtel und gewöhnliche Bettlaken werden in Verwahrung genommen. Geeignete Schutzkleidung und geeignetes Bettmaterial sind bereitzustellen. Die Beobachtung wird entsprechend der ärztlich zu präzisierenden Schutzanforderung organisiert. Der medizinische Dienst wurde verständigt. Über Art und Fortdauer jeder Maßnahme wird gesondert entschieden.

2. Anlass

Frau Finkbeiner befindet sich in einer akuten Krise und äußert eine ernst zu nehmende Gefährdung ihrer eigenen Sicherheit. Der Dienst geht von unmittelbarem Handlungsbedarf aus. Die Einzelheiten der medizinischen Anamnese werden in der Gesundheitsakte geführt und nicht in diesen allgemein zugänglichen Verwaltungsauszug aufgenommen. Die Entscheidung ist eine Schutzmaßnahme und keine Disziplinarmaßnahme.

Die Anstaltsleitung stützt die besonderen Sicherungsmaßnahmen auf die einschlägigen Regelungen des JVollzGB Buch 3, insbesondere § 67. Wegen der unverzüglichen Reaktion konnte eine vollständige ärztliche Stellungnahme vor dem Vollzug nicht abgewartet werden. Die ärztliche Beteiligung wird umgehend nachgeholt. Das Ergebnis, mögliche Änderungen und die Ausgabe von Ersatzmaterial sind zu dokumentieren. Ein späterer ruhiger Eindruck ersetzt keine fachliche Bewertung.

Die Gefangene soll in verständlicher Sprache erfahren, welche Maßnahme warum besteht und wann eine neue Beurteilung stattfindet. Annemarie Seibold, stellvertretende Anstaltsleiterin.''')
    doc(c,10,'pdf','Bekanntgabe_und_Verwahrung', '''JVA Silberrain. Ausgabevermerk SR-227/26 vom 29.09.2026.

Die Anordnung wurde Frau Finkbeiner am 28.09. um 19.00 Uhr erklärt. Wegen ihres Zustands wurde der schriftliche Text am 29.09. um 09.15 Uhr nochmals vollständig ausgehändigt und mit ihr durchgesprochen. Sie wollte zu diesem Zeitpunkt keine Empfangsunterschrift leisten. Anwesend waren Daria Mertens und Hannelore Kern. Ob die erste mündliche Bekanntgabe inhaltlich vollständig erfasst wurde, lässt sich aus diesem Vermerk nicht ableiten.

Verwahrt wurden ein Gürtel und die gewöhnlichen Bettlaken. Im Materialjournal ist am 28.09. um 19.30 Uhr die Ausgabe einer Schutzdecke und eines Satzes geeigneter Schutzkleidung eingetragen. Eine zweite Decke wurde am 29.09. um 07.10 Uhr nachgefordert und um 07.35 Uhr ausgegeben. Der Grund lautet „Kälteempfinden, zusätzliche Abdeckung“. Die Liste verzeichnet keine Temperaturmessung im Raum.

Die verwahrten Gegenstände sind weiterhin Eigentum der Gefangenen. Eine Einziehung oder Vernichtung wurde nicht angeordnet. Rückgabe und Ersatzversorgung sind nach dem jeweils aktuellen Schutzbedarf zu überprüfen. Die Verwahrungsliste enthält noch keine spätere Entscheidung über den 01.10. hinaus.

Hannelore Kern, Vollzugsdienst. Die Dokumentation beschreibt keine Selbstschädigungsmethode.''')
    doc(c,11,'pdf','Beobachtungsnotiz_29_September', '''JVA Silberrain, geschützter Bereich. Auszug aus dem Dienstbericht vom 29.09.2026. Daria Mertens, Dienst von 06.00 bis 14.00 Uhr.

Um 07.05 Uhr meldete Frau Finkbeiner, dass ihr kalt sei. Eine zweite geeignete Decke wurde angefordert. Die Ausgabe ist um 07.35 Uhr bestätigt. Sie trank Tee und nahm das Frühstück teilweise an. Um 08.20 Uhr führte eine Ärztin ein Gespräch mit ihr. Die dienstliche Notiz enthält keine Diagnose. Eine eigene Bewertung, dass die Gefahr vollständig entfallen sei, wurde vom Vollzugsdienst nicht vorgenommen.

Frau Finkbeiner fragte wiederholt, wann sie den Raum verlassen könne. Ich erklärte, dass die Entscheidung nach der fachlichen Beurteilung getroffen werde. Im Nachhinein kann ich nicht sagen, ob ihr ein konkreter Zeitpunkt genannt wurde. Die Nachfrage zur Privatsphäre habe ich an die Dienstleitung weitergegeben. Sie wollte eine Erklärung durch eine Mitarbeiterin und keine Besprechung vor mehreren Personen.

Im Übergabefeld steht „ruhig, kooperativ“. Das bezieht sich auf mein beobachtetes Verhalten während dieser Schicht. Es bedeutet weder eine medizinische Entwarnung noch eine Bewertung der vorausgegangenen Nacht. Die nächtliche Beobachtungsliste ist in diesem Auszug nicht enthalten. Für eine vollständige Rekonstruktion müssen beide Listen zusammengeführt werden.

Die Gefangene berichtet am Ende der Schicht, sie habe sich bei einer Begleitung am Vorabend grob angesprochen gefühlt. Sie konnte in diesem Gespräch keine Person sicher benennen. Ich habe angeboten, die Beschwerde in Ruhe aufzunehmen. Ein Gespräch mit der Leitung ist vorgemerkt.''')
    doc(c,12,'pdf','Belegungsmitteilung_06_Oktober', '''JVA Silberrain, Belegungsdienst. Vermerk vom 06.10.2026, 11.30 Uhr, Raum B 12.

Frau Finkbeiner zog am 01.10. um 10.20 Uhr aus dem besonders gesicherten Bereich in B 12. Der Raum war zu diesem Zeitpunkt mit einer weiteren Gefangenen belegt. Beide Frauen wurden zuvor angesprochen. Die Liste enthält jeweils das Kürzel „einverstanden“. Eine gesonderte schriftliche Erklärung der Mandantin über Umfang und Dauer dieses Einverständnisses liegt diesem Auszug nicht bei.

Wegen eines Wasserschadens im benachbarten Raum wurde am 03.10. um 18.00 Uhr vorübergehend eine dritte Frau aufgenommen. Sie verließ B 12 am 06.10. um 09.00 Uhr. Eine vorherige schriftliche Zustimmung von Frau Finkbeiner zur dritten Person ist nicht verzeichnet. Die andere ursprüngliche Bewohnerin äußerte gegenüber dem Dienst, sie wolle die Lage für wenige Tage mittragen. Daraus folgt keine Erklärung für Frau Finkbeiner.

Die vorübergehende Unterbringung wurde mit dem besonderen organisatorischen Bedarf und der begrenzten Zahl geeigneter Plätze begründet. Der Raum B 15 war wegen einer technischen Prüfung gesperrt. Ein genauer Termin für die Freigabe lag am 03.10. noch nicht vor. Die Belegung war als Übergangslösung gedacht; über konkrete tägliche Öffnungszeiten sagt dieser Vermerk nichts aus.

Belegungsdienst Paula Wetter. Namen und Daten der beiden weiteren Frauen sind in dieser Arbeitskopie gekürzt, da ihre persönlichen Verhältnisse nicht Gegenstand des Mandats sind.''')
    doc(c,13,'pdf','Aerztliche_Mitteilung_14_Oktober', '''Medizinischer Dienst der JVA Silberrain. Mitteilung vom 14.10.2026 an die bevollmächtigte Rechtsvertretung. Verfasserin Dr. Merle Auer.

Am Abend des 28.09. wurde der ärztliche Bereitschaftsdienst über eine akute Krise informiert. Eine erste telefonische Einschätzung und ein persönlicher Kontakt sind in der Gesundheitsakte dokumentiert. Ein ausführlicheres Gespräch fand am Morgen des 29.09. statt. Der Vermerk des Vollzugsdienstes ersetzt die ärztliche Dokumentation dieser Kontakte nicht. Die vollständige Behandlungsakte wird nach Abgleich der Einwilligung auf sicherem Weg gesondert übersandt.

Der Schutzbedarf wurde zunächst als erheblich eingeschätzt. Der spätere ruhigere Eindruck ließ eine stufenweise Veränderung zu, bedeutete aber keine pauschale Entwarnung für jeden Gegenstand und jede Umgebung. Am 01.10. wurde der Wechsel in einen geeigneten normalen Haftraum mit fortgesetzten Gesprächen unterstützt. Die Frage, ob eine Mitbelegung im konkreten Fall hilfreich oder belastend wirkt, sollte mit beiden Betroffenen und dem psychologischen Dienst abgestimmt werden.

Die Rückgabe gewöhnlicher Bettwäsche wurde am 04.10. befürwortet. Über den Gürtel war zunächst nach einer weiteren Beurteilung zu entscheiden. In der übermittelten Verwaltungskopie fehlt der Eintrag über die Überprüfung vom 07.10. Daraus lässt sich nicht sicher schließen, dass sie unterblieben ist; die vollständige Dokumentation muss beigezogen werden.

Aktuell ist regelmäßige Betreuung vorgesehen. Eine akute Veränderung muss sofort fachlich bewertet werden. Diese Mitteilung enthält keine Anweisung zur Selbstbehandlung und ersetzt keine laufende medizinische Betreuung. Sie entscheidet nicht über die Rechtmäßigkeit der gerichteten Anträge.''')
    doc(c,14,'pdf','Brief_Idris_Finkbeiner', '''Idris Finkbeiner, Fiktive Hopfenzeile 17, 72108 Rottenburg. An Amira Finkbeiner. 08.10.2026.

Liebe Amira,

Nora hat die Kastanien diesmal nach Größe sortiert und behauptet, das sei schon Mathematik. Ich soll Dir unbedingt sagen, dass die kleine mit dem weißen Fleck für Dich ist. Ich habe ihr erklärt, dass es Dir nicht gut ging und dass Menschen da sind, die Dir helfen. Mehr wollte sie heute nicht wissen.

Am Telefon am Samstag klangst Du ruhiger. Ich habe der Anwältin deshalb gesagt, ich glaube, es gehe Dir etwas besser. Ich kann natürlich nicht beurteilen, ob Du wieder alles zurückbekommen solltest. Das war keine ärztliche Einschätzung. Wenn im Schreiben etwas anderes daraus geworden ist, muss ich das richtigstellen.

Die Raumgröße habe ich nur aus Deiner Beschreibung geschätzt. Ich war nie in B 12. Bei zwölf Quadratmetern dachte ich an unser altes Wohnzimmer, aber das ist kein verlässlicher Vergleich. Du hattest mir gesagt, dass vor allem nachts keine Ruhe sei und dass Du Dich beim Waschen nicht gut abschirmen könntest. Ich hätte das wohl genauso weitergeben sollen, statt eine Zahl zu nennen.

Ich bringe zum Besuch die erlaubten Unterlagen mit und frage vorher nach. Bitte sag mir, welche persönlichen Dinge ich der Kanzlei erzählen darf. Ich möchte helfen, ohne Dir das Gefühl zu geben, dass alle über Dich sprechen. Dein Idris.''')
    doc(c,15,'pdf','Folgeverfuegung_01_Oktober', '''JVA Silberrain. SR-227/26. Verfügung vom 01.10.2026, 09.40 Uhr.

1. Änderung der Schutzmaßnahmen

Die Unterbringung im besonders gesicherten Haftraum endet heute mit der Verlegung nach B 12. Der Vollzug erfolgt um 10.20 Uhr. Die weitere Betreuung und die angepasste Beobachtung werden nach der ärztlichen Empfehlung organisiert. Die ursprüngliche besonders gesicherte Unterbringung wird damit nicht stillschweigend fortgesetzt.

Die gewöhnliche Bettwäsche bleibt zunächst verwahrt; geeignetes Bettmaterial wird in B 12 bereitgestellt. Die Ausgabe gewöhnlicher Bettwäsche wird nach erneuter fachlicher Rückmeldung überprüft. Der Gürtel bleibt vorerst in Verwahrung. Für beide Gegenstände sind die Gründe und das Ergebnis der nächsten Überprüfung getrennt zu dokumentieren. Eine unbefristete pauschale Rückgabesperre wird nicht angeordnet.

2. Gemeinsame Unterbringung

Die Mandantin soll zunächst mit einer weiteren Gefangenen in B 12 leben. Beide Personen sind anzuhören. Die Rückmeldung des psychologischen Dienstes wird in die weitere Planung einbezogen. Bei erheblichen Spannungen ist die Abteilung unverzüglich zu informieren. Eine Gemeinschaftsunterbringung ersetzt keine medizinische Betreuung und macht die Mitgefangene nicht für die Sicherheit von Frau Finkbeiner verantwortlich.

Nachtrag vom 04.10., 13.20 Uhr: Gewöhnliche Bettwäsche wurde nach medizinischer Rückmeldung ausgegeben. Über den Gürtel steht eine weitere dokumentierte Entscheidung aus. Annemarie Seibold.''')
    doc(c,16,'pdf','Beschluss_09_Oktober', '''Fiktives Gerichtsdokument. Strafvollstreckungskammer. Zeichen ausschließlich für diese Arbeitsakte: ÜBUNG 4 StVK 307/26. Beschluss vom 09.10.2026.

1. Tenor

Der Antrag auf vorläufige Aufhebung sämtlicher Schutzmaßnahmen und Zuweisung einer Einzelzelle wird abgelehnt. Über den Hauptsacheantrag wird nach Eingang weiterer Unterlagen entschieden.

2. Gründe

Nach Mitteilung der Anstalt bestand eine erhebliche Eigengefährdung. Die Gefangene wurde inzwischen aus dem besonders gesicherten Bereich verlegt. Eine aktuelle fachliche Stellungnahme der Antragstellerseite liegt nicht vor. Der Ehemann kann die Gefährdung nicht medizinisch einschätzen. Die Kammer sieht deshalb keinen Anlass, der Anstalt vorläufig die Rückgabe sämtlicher Gegenstände aufzugeben.

Die Anstalt gibt für B 12 eine Gesamtfläche von 14,20 Quadratmetern an. Bei der inzwischen wieder bestehenden Doppelbelegung errechnen sich rechnerisch 7,10 Quadratmeter pro Person. Die Kammer hält diese Angabe für ausreichend, um eine dringliche unzumutbare Enge im vorliegenden Eilverfahren zu verneinen. Den abweichenden Planstand, den Anteil der Sanitärzone und die vorübergehende Dreifachbelegung hat sie noch nicht näher aufgeklärt. Die genauen Einschlusszeiten bleiben dem Hauptsacheverfahren vorbehalten.

Die begehrte Einzelunterbringung und vollständige Rückgabe würde wesentliche Teile des Hauptbegehrens vorwegnehmen. Eine weitere vorläufige Regelung wird nicht getroffen. Das Gericht hat die vollständige ärztliche Dokumentation noch nicht beigezogen. Eine Entscheidung über Schadensersatz oder eine Feststellung eines Übergriffs ist nicht Gegenstand dieses Beschlusses.

Dr. Gundula Ehrenfeld, Richterin. Die Entscheidung ist erfunden und keine zitierbare Rechtsprechungsquelle.''')
    doc(c,17,'pdf','Postjournal_und_Empfang', '''Kanzlei Kappel. Fiktiver Auszug aus dem Postjournal vom 16.10.2026. Vorgang SK 155/26.

Am 02.10. um 15.02 Uhr ging der Schriftsatz mit der Überschrift „Beschwerde“ elektronisch bei der Strafvollstreckungskammer ein. Die technische Eingangsbestätigung nennt zwei Anlagen. Am 05.10. teilte die Geschäftsstelle mit, der Vorgang werde als Hauptsacheantrag und vorläufiges Begehren unter ÜBUNG 4 StVK 307/26 geführt. Eine andere Gerichtszuständigkeit wurde damit nicht festgestellt.

Der Beschluss vom 09.10. ging am 12.10. um 08.51 Uhr in der Kanzlei ein. Rechtsanwältin Kappel gab das elektronische Empfangsbekenntnis mit Zustelldatum 12.10.2026 ab. Die Mandantin erhielt die weitergeleitete Kopie nach Stationsvermerk am 14.10. Die Daten sind getrennt zu erfassen.

Am 13.10. um 17.08 Uhr wurde ein Schriftsatz „Sofortige Beschwerde“ übermittelt. Er beantragt die Aufhebung des Beschlusses vom 09.10. und verweist auf die erste Eingabe. Weitere Tatsachen oder Anlagen sind nicht beigefügt. Ein Vermerk zur Statthaftigkeit liegt in der Papierakte nicht. Eine Entscheidung über diese Eingabe wurde bis zum 16.10. um 08.00 Uhr nicht empfangen.

Der vorliegende Ausdruck berechnet keine Frist und bestätigt keine Zulässigkeit. Er enthält keine echten elektronischen Postfachkennungen oder Signaturdaten.''')
    doc(c,18,'pdf','Interner_Ablauf_Schutzbereich', '''JVA Silberrain. Arbeitsanweisung der Abteilung B, Auszug vom 01.07.2026, übermittelt am 07.10.2026. Kein amtlicher Gesetzes- oder Verordnungstext.

Über die Ausgabe geeigneter Ersatzkleidung und geeigneten Bettmaterials ist eine Materialliste zu führen. Beschwerden über Kälte, Unverträglichkeit oder unzureichenden Sichtschutz werden an die Dienstleitung und gegebenenfalls an den medizinischen Dienst weitergegeben. Die bloße Verwahrung eines Gegenstands ist kein Grund, notwendige Ersatzversorgung zu unterlassen.

Anordnung, tatsächlicher Vollzug, spätere Änderung und Beendigung besonderer Sicherungsmaßnahmen sind getrennt einzutragen. Ein allgemeiner Vermerk „ruhig“ ersetzt keine fachliche Entscheidung. Für den besonders gesicherten Bereich werden die ärztlichen Kontakte in der Gesundheitsakte und die organisatorischen Kontakte im Stationsjournal dokumentiert. Eine Arbeitskopie kann deshalb unvollständig sein.

Vor gemeinsamer Unterbringung sind mögliche schädliche Einflüsse, Zustimmung soweit erforderlich und die Situation aller Betroffenen zu prüfen. Eine Mitgefangene übernimmt keine medizinische Überwachungspflicht. Bei zeitweiser Mehrfachbelegung aus organisatorischem Anlass dokumentiert der Belegungsdienst Beginn, Ende und Alternativen. Eine Hausanweisung schafft keine von den gesetzlichen Voraussetzungen unabhängige Eingriffsbefugnis.

Die Arbeitsanweisung enthält keine feste Quadratmetergrenze. Raumpläne und tatsächliche Belegung sind zusammen mit den weiteren Unterbringungsbedingungen zu betrachten.''')
    doc(c,19,'pdf','Rechnung_Kappel_10_Oktober', '''Kanzlei Kappel, Fiktiver Mühlrain 13, 72070 Tübingen. Rechnung SK-2026-87 vom 10.10.2026 an Amira Finkbeiner.

Für die vereinbarte erste Beratung und außergerichtliche Korrespondenz wird ein Nettohonorar von 300,00 EUR berechnet. Hinzu kommen 19 Prozent Umsatzsteuer, 57,00 EUR. Der Gesamtbetrag beträgt 357,00 EUR. Die Vereinbarung datiert vom 01.10.2026 und soll mit der übrigen Mandatsakte übergeben werden. Die vorliegende Rechnung ist keine vollständige gerichtliche Gebührenabrechnung.

Zahlungsstand am 16.10.2026: offen. Idris Finkbeiner hat um eine kurze Erläuterung gebeten, welche Tätigkeiten erfasst sind und ob ein weiterer Betrag zu erwarten ist. Eine Zahlungszusage des Ehemanns ist nicht dokumentiert. Er ist nicht automatisch Schuldner der Rechnung, nur weil er die organisatorische Kommunikation übernommen hat.

Auf der Rechnung ist für diese Übungsakte bewusst keine echte Bankverbindung abgebildet. Die Kostenfrage ist von der dringlichen Schutz- und Unterbringungsprüfung zu trennen. Eine Erstattung durch die Staatskasse oder die Anstalt wird nicht zugesagt.''')
    doc(c,20,'pdf','Raumblatt_B12_und_Aufmass', '''JVA Silberrain, technischer Dienst. Raum B 12. Zusammenstellung vom 07.10.2026 durch Benno Stadel.

1. Bestandsplan

Der Bestandsplan von 2018 nennt 4,00 m mal 3,55 m und damit 14,20 m² Gesamtfläche einschließlich einer mit 1,80 m² bezeichneten Sanitärzone. Eine Änderung der Installationswand ist darin nicht eingetragen. Der Ausdruck enthält keine Bestätigung, dass diese Maße nach der Reparatur 2024 unverändert geblieben sind.

2. Aufmaß vom 07.10.2026

Das heutige Aufmaß zwischen den Innenflächen ergibt 3,96 m mal 3,50 m und damit 13,86 m². Die Sanitärzone einschließlich Wandanteil wurde mit 1,92 m² angesetzt. Der verbleibende rechnerische Bereich beträgt 11,94 m². Das ist eine technische Flächenaufteilung, keine rechtliche Festlegung einer Mindestfläche. Feste Möbel, Verkehrsflächen und nutzbare Bereiche sind gesondert zu betrachten.

Die Sanitärtür schließt nach Sichtprüfung nicht vollständig; der Spalt an der Schlossseite beträgt ungefähr zwei Zentimeter. Der Reparaturauftrag wurde am selben Tag angelegt. Ob die Tür während der Dreifachbelegung in gleicher Weise stand, ist nicht dokumentiert. Das Fenster lässt sich nur begrenzt öffnen. Luftmessungen wurden bei diesem Termin nicht vorgenommen.

3. Einrichtung und Nutzung

Im Raum stehen zwei reguläre Betten, ein vorübergehend aufgestelltes drittes Bett, ein Tisch und drei Stühle. Das dritte Bett wurde am 06.10. wieder entfernt. Die vorhandene Skizze zeigt den Zustand nach der Entfernung. Für den Zeitraum davor sind Fotos angefordert; sie liegen dieser Akte noch nicht bei. Eine schematische Zeichnung ist dem Aufmaß beigefügt, aber nicht maßstabsgerecht. Aus der reinen Flächenzahl lassen sich die täglichen Einschlusszeiten nicht ableiten.''')
    return c

def populate_mails(c):
    if c['land']=='Hessen':
        rows=[
        (21,'2026-09-19','Marlene Hagedorn','marlene.hagedorn@postfach.example','v.brack@kanzlei-brack.example','Unterlagen meines Bruders', '''Sehr geehrter Herr Brack, anbei ist die Kopie der Anordnung, die mein Bruder mir geschickt hat. Ich habe sie erst heute aus dem Briefkasten genommen. Ich habe keine Vollmacht, für ihn Anträge zurückzunehmen oder zuzustimmen. Er hat mich gebeten, Ihnen bei der Organisation zu helfen. Mein Vater ist im Krankenhaus gewesen; der Brief darüber kam bei Jannis wohl sehr spät an. Das entschuldigt keinen Streit, erklärt aber, warum ihn das beschäftigt. Bitte sagen Sie mir, ob Sie mit ihm persönlich sprechen können. Ich weiß weder, wie lange die Trennung noch dauern soll, noch was im ärztlichen Dienst passiert ist. Die Kopie hat eine schlecht lesbare Rückseite; ich schicke nichts weg, bevor die Originale gesichert sind. Freundliche Grüße, Marlene Hagedorn.''',[9]),
        (22,'2026-09-25','Volker Brack','v.brack@kanzlei-brack.example','marlene.hagedorn@postfach.example','Beschluss im vorläufigen Verfahren', '''Sehr geehrte Frau Hagedorn, das Gericht hat das vorläufige Begehren abgelehnt. Den Beschluss füge ich bei. Die Hauptsache ist noch nicht entschieden. Ich habe Ihrem Bruder eine Kopie per Post geschickt. Bitte verstehen Sie die Nachricht nicht als Aufforderung, selbst einen Rechtsbehelf einzulegen. Die gerichtliche Entscheidung bezieht sich auf den bisherigen Vortrag; die Anstaltsakte ist noch nicht vollständig in unserer Kanzlei. Ich bin nächste Woche auf Fortbildung und bitte meine Mitarbeiterin, eingehende Unterlagen sofort vorzulegen. Den genauen weiteren Schritt möchte ich danach mit Ihrem Bruder besprechen. Eine neue Sachentscheidung der Anstalt oder ein Untersuchungsergebnis liegt mir heute nicht vor. Mit freundlichen Grüßen, Volker Brack.''',[16]),
        (23,'2026-10-12','Dr. Ottilie Gründel','vollzug@jva-waldhain.example',c['lawmail'],'Fortdauer und fehlender Empfangseintrag', '''Sehr geehrte Frau Rehbein, die Vertretungsanzeige ist eingegangen. Wir übersenden die Fortdauerentscheidung. Der Ausgabevermerk ist in der vorliegenden Kopie unvollständig; die Station prüft den genauen Zeitpunkt. Bitte verwenden Sie das Datum der E-Mail nicht ohne Weiteres als Datum einer Bekanntgabe an Herrn Hagedorn. Die Rückmeldung der Aufsicht wurde intern angefordert. Zu den geplanten Kontakten sind mehrere Dienste zu beteiligen. Die Anstalt hält die ursprüngliche Lage für ernst, schließt einen stufenweisen anderen Ablauf aber nicht aus. Wir können Ihnen heute noch keinen abgeschlossenen Plan zusenden. Mit freundlichen Grüßen, Ottilie Gründel.''',[15]),
        (24,'2026-10-14','Mira Eschenbach','medizin@jva-waldhain.example',c['lawmail'],'Teilakte und gesonderte Einwilligung', '''Sehr geehrte Frau Rehbein, Ihre gezielte Einwilligung ist angekommen. Anbei finden Sie zunächst die Mitteilung vom 08.10. Die vollständigen Befundunterlagen müssen aus dem medizinischen System zusammengestellt werden. Ein Eintrag in einer Dienstübergabe ist keine vollständige ärztliche Dokumentation. Wir werden insbesondere prüfen, welche Kontakte am 10. und 12.09. erfolgten und welche Unterlagen den Beschwerden vom 30.09. zugeordnet sind. Die Gesundheitsakte wird nicht an Angehörige verteilt. Wir bitten um einen direkten sicheren Übermittlungsweg. Bei aktuellen Beschwerden veranlasst die Station unabhängig davon die medizinische Rückmeldung. Mit freundlichen Grüßen, Mira Eschenbach.''',[13]),
        (25,'2026-10-13','Finja Mohn','sozialdienst@jva-waldhain.example',c['lawmail'],'Kontaktangebot außerhalb der bisherigen Gruppe', '''Sehr geehrte Frau Rehbein, ich habe Herrn Hagedorn am 17.09. und kurz am 06.10. gesprochen. Er möchte nicht zu den beiden am Streit beteiligten Personen zurück. Ich habe ein begleitetes Gespräch mit einem anderen Gefangenen aus der Werkgruppe vorgeschlagen. Ein geeigneter Raum ist grundsätzlich vorhanden, die Begleitung wurde noch nicht abschließend zugesagt. Mein Vorschlag ist keine Aufhebung der Sicherungsanordnung. Die Kontaktliste enthält außerdem mehrere Türgespräche, deren Dauer nicht überall eingetragen ist. Ich bitte darum, diese nicht als längere geplante Gespräche zusammenzuzählen. Herr Hagedorn war am 06.10. verärgert, blieb aber im Gespräch. Eine fachliche Gefahrenbewertung nehme ich hiermit nicht vor. Freundliche Grüße, Finja Mohn.''',[]),
        (26,'2026-10-15','Gero Tann','technik@jva-waldhain.example',c['lawmail'],'Raumblatt S 4', '''Sehr geehrte Frau Rehbein, beigefügt ist das Raumblatt mit der Übergabenotiz. Der Bestandsplan stammt aus der Zeit vor einer Leitungsreparatur. Es gibt deshalb keinen Grund, die Fläche auf den Zentimeter als neu vermessen zu behandeln. Die Sanitärzone ist in der Gesamtfläche enthalten. Aussagen über tatsächliche Einschlusszeiten oder Kontakte kann der technische Dienst nicht treffen. Diese Daten führt die Abteilung. Wenn Sie ein neues Aufmaß benötigen, richten Sie die Anfrage bitte über die Vollzugsleitung; wir vereinbaren dann einen Termin, bei dem die Nutzung möglichst wenig beeinträchtigt wird. Mit freundlichen Grüßen, Gero Tann.''',[20]),
        (27,'2026-10-14','Levin Bering','stationsleitung@jva-waldhain.example',c['lawmail'],'Ergänzung zur Übergabe vom September', '''Sehr geehrte Frau Rehbein, ich antworte nach Abstimmung mit der Abteilungsleitung. Mein Eintrag lautete „Arzt informiert“. Ich habe damit keine abgeschlossene Untersuchung bestätigt. Der Kollege hat in einer anderen Liste „medizinisch abgeklärt“ notiert; ich kann nicht sagen, worauf er diese Formulierung stützte. Meinen Bericht habe ich als Kopie zur Abteilungsakte gegeben. Ich bitte, die unterschiedlichen Formulierungen nicht in eine neue einheitliche Fassung umzuschreiben. Es ist richtig, dass ich nicht während des gesamten Einsatzes freie Sicht hatte. Zu einem Zeitraum, den ich nicht gesehen habe, kann ich keine sichere Aussage treffen. Mit freundlichen Grüßen, Levin Bering.''',[11]),
        (28,'2026-10-15',c['lawyer'],c['lawmail'],'vollzug@jva-waldhain.example','Bitte um Sicherung der vorhandenen Unterlagen', '''Sehr geehrte Frau Dr. Gründel, bitte sichern Sie die unveränderten Dienstberichte, die zugehörigen Änderungsvermerke, den vollständigen Ablauf der ärztlichen Befassung und die Unterlagen zur aufsichtsbehördlichen Beteiligung. Ich bitte um Mitteilung, ob zu dem Vorfall vom 09.09. überhaupt Bildaufzeichnungen bestanden und gegebenenfalls welche regulären Speicherfristen gelten. Das Vorhandensein einer Aufnahme wird nicht unterstellt. Bitte erstellen Sie keine nachträgliche gemeinsame Erinnerung anstelle der einzelnen Berichte. Für die weitere Bearbeitung benötigen wir die Quellen in ihrer jeweiligen Fassung. Die Anfrage enthält weder einen Verzicht auf Rechtsschutz noch einen abschließend erhobenen strafrechtlichen Vorwurf. Mit freundlichen Grüßen, Malin Rehbein.''',[]),
        (29,'2026-10-16','Hiltrud Bär','vollzug@jva-waldhain.example',c['lawmail'],'Eingang Ihrer Sicherungsbitte', '''Sehr geehrte Frau Rehbein, Ihre Bitte wurde heute um 08.15 Uhr an die zuständigen Stellen weitergegeben. Im Gemeinschaftsraum bestand nach Auskunft der Technik keine Bildaufzeichnung. Ob der vorgelagerte Flur in dem fraglichen Zeitraum erfasst wurde und ob Daten noch vorhanden sind, wird geprüft. Wir können eine Sicherung noch nicht als abgeschlossen bestätigen. Die schriftlichen Dienstberichte bleiben unverändert in der Akte; Ergänzungen werden als solche datiert. Die medizinischen Unterlagen werden getrennt bearbeitet. Bitte bezeichnen Sie bei weiterer Korrespondenz den jeweiligen Vorgang, damit Sicherungsanordnung, körperliche Beschwerden und Ihr Auskunftsantrag nicht miteinander verwechselt werden. Mit freundlichen Grüßen, Hiltrud Bär.''',[]),
        (30,'2026-10-16','Nika Lenz','assistenz@kanzlei-rehbein.example',c['lawmail'],'Vollmacht und Termin mit Herrn Hagedorn', '''Frau Rehbein, die unterschriebene Vollmacht liegt vor; die abgetippte Arbeitskopie ist beigefügt. Die Besuchsabteilung bietet am 19.10. um 13.00 Uhr einen anwaltlichen Besuch an. Die Reservierung ist noch nicht bestätigt. Ich habe den Termin als organisatorische Wiedervorlage eingetragen, nicht als gerichtliche Frist. In der Liste fehlen die Bekanntgabe der Fortdauerentscheidung und der aktuelle Hauptsachestand. Die Seitenzahlen der Anlagen dürfen wir erst nach dem endgültigen Zusammenstellen vergeben. Bitte sagen Sie mir, ob ich für das Gespräch die beiden unterschiedlichen Dienstberichte nebeneinander ausdrucke. Nika Lenz.''',[2]),
        (31,'2026-10-15','Marlene Hagedorn','marlene.hagedorn@postfach.example',c['lawmail'],'Korrektur zu unserem Telefonat', '''Sehr geehrte Frau Rehbein, ich habe meinen Kalender angesehen. Das Telefonat mit meinem Bruder war am Dienstag, 29.09., nicht am Montag. Ich hatte Ihnen das am Telefon falsch gesagt. Anbei ist der Brief, den ich danach geschrieben habe. Ich möchte nicht, dass die Situation dramatischer klingt, als ich sie selbst kenne. Ich habe die Schulter nicht gesehen. Jannis sagte, dass ihm manche Bewegungen wehtun. Mein Vater ist inzwischen wieder zu Hause, deshalb ist seine Sorge etwas kleiner geworden. Bitte richten Sie meinem Bruder aus, dass ich einen Besuch auch unter den üblichen Kontrollen wahrnehmen würde. Freundliche Grüße, Marlene Hagedorn.''',[14]),
        (32,'2026-10-16','Volker Brack','v.brack@kanzlei-brack.example',c['lawmail'],'Übergabe und ergänzende Dateiliste', '''Sehr geehrte Kollegin Rehbein, beigefügt übersende ich den Übergabevermerk. Ihre Mitarbeiterin hat mir die Chronologie zur Vollständigkeitskontrolle geschickt; ich gebe sie unverändert mit zurück. Ich bestätige darin nur die von meiner Kanzlei selbst versandten und empfangenen Vorgänge. Den Tagesablauf in der JVA kann ich nicht bestätigen. Bitte beachten Sie meine fehlende Prüfung zum Rechtsbehelf und den Umstand, dass die ärztliche Stellungnahme nicht vollständig vorliegt. Die Hauptsache habe ich nicht zurückgenommen. Sollte Ihnen ein Sendeprotokoll fehlen, nennen Sie bitte das Datum; die Mitarbeiterin kann es gesondert aus dem Journal bereitstellen. Mit kollegialen Grüßen, Volker Brack.''',[7,33])]
    else:
        rows=[
        (21,'2026-10-01','Idris Finkbeiner','idris.finkbeiner@postfach.example','s.kappel@kanzlei-kappel.example','Bitte um Hilfe für meine Frau', '''Sehr geehrte Frau Kappel, meine Frau möchte anwaltliche Unterstützung. Sie hat mir am Telefon gesagt, dass sie bald in einen anderen Raum komme. Ich weiß nicht, welche Gegenstände sie dann zurückbekommt. Bitte besprechen Sie das mit ihr und nicht nur mit mir. Sie möchte nicht, dass unsere Tochter medizinische Einzelheiten erfährt. Ich habe keine Unterlagen über die genaue Gefahr und möchte diese auch nicht selbst bewerten. Heute klang sie ruhiger. Das ist mein persönlicher Eindruck nach einem kurzen Gespräch. Wenn Sie eine Vollmacht brauchen, bitte ich um den für die JVA vorgesehenen Weg. Freundliche Grüße, Idris Finkbeiner.''',[]),
        (22,'2026-10-07','Benno Stadel','technik@jva-silberrain.example','s.kappel@kanzlei-kappel.example','Aufmaß B 12 und älterer Planstand', '''Sehr geehrte Frau Kappel, anbei erhalten Sie unsere Zusammenstellung. Die Zahl 14,20 stammt aus dem Bestandsplan, die Zahl 13,86 aus dem heutigen Aufmaß. Beide enthalten zunächst die Sanitärzone. Bitte verwenden Sie nicht eine Zahl aus der einen Fassung und einen Abzug aus der anderen, ohne das zu erläutern. Das dritte Bett war beim Aufmaß bereits entfernt. Die Tür zur Sanitärzone wurde beanstandet und ein Reparaturauftrag angelegt. Über den tatsächlichen Zustand an früheren Tagen kann ich nur nach den Dienstnotizen Auskunft geben. Der technische Dienst bewertet keine rechtliche Mindestfläche. Mit freundlichen Grüßen, Benno Stadel.''',[20]),
        (23,'2026-10-06','Paula Wetter','belegung@jva-silberrain.example','s.kappel@kanzlei-kappel.example','Zeitweilige dritte Belegung', '''Sehr geehrte Frau Kappel, die dritte Frau hat B 12 heute Morgen verlassen. Den Ablauf finden Sie in der Anlage. Ursache war der Wasserschaden im benachbarten Bereich. Für die ursprüngliche Doppelbelegung enthält die Liste ein Einverständniskürzel, für die dritte Person nicht. Wir können aus dem Kürzel nicht nachträglich eine umfassendere Zustimmung erstellen. Die technische Sperrung von B 15 wird gesondert dokumentiert. Wenn Sie eine weitergehende Auskunft zu Alternativen benötigen, leiten wir die Anfrage an die zuständige Leitung weiter. Die Gesundheitsdaten der anderen Frauen werden nicht ungeprüft mitgegeben. Mit freundlichen Grüßen, Paula Wetter.''',[12]),
        (24,'2026-10-14','Dr. Merle Auer','medizin@jva-silberrain.example',c['lawmail'],'Mitteilung und vollständige Behandlungsunterlagen', '''Sehr geehrter Herr Lindenfels, die ergänzte Einwilligung ist eingegangen. Die beigefügte Mitteilung dient der ersten Orientierung. Die vollständige Behandlungsakte folgt gesondert auf sicherem Weg. Bitte trennen Sie die Befürwortung einer Verlegung aus dem besonders gesicherten Bereich von einer Aussage über jeden einzelnen Gegenstand. Wir haben keine generelle Entwarnung für jede Situation erklärt. Die regelmäßigen Gespräche laufen weiter. Eine mögliche aktuelle Verschlechterung wird unabhängig vom gerichtlichen Verfahren sofort behandelt. Ihre Bitte um den vollständigen Eintrag vom 07.10. wurde an die Dokumentationsstelle gegeben. Mit freundlichen Grüßen, Merle Auer.''',[13]),
        (25,'2026-10-12','Sigrun Kappel','s.kappel@kanzlei-kappel.example','idris.finkbeiner@postfach.example','Entscheidung zum Eilverfahren', '''Sehr geehrter Herr Finkbeiner, das Gericht hat den vorläufigen Antrag abgelehnt. Den Beschluss füge ich bei. Der Hauptsacheantrag bleibt offen. Bitte leiten Sie daraus nicht ab, die gesamte bisherige Behandlung sei abschließend rechtmäßig festgestellt. Die Kammer hat bestimmte Tatsachen noch nicht aufgeklärt. Ich werde den weiteren Schritt mit Ihrer Frau besprechen. Sie haben gefragt, ob die drei Tage mit einer weiteren Person nun „weg“ seien. Nein, der Zeitraum bleibt Teil der Akte. Welche gerichtliche Entscheidung dazu beantragt werden kann, muss gesondert geprüft werden. Mit freundlichen Grüßen, Sigrun Kappel.''',[16]),
        (26,'2026-10-13','Idris Finkbeiner','idris.finkbeiner@postfach.example',c['lawmail'],'Meine Schätzung war keine Messung', '''Sehr geehrter Herr Lindenfels, ich möchte meine frühere Angabe richtigstellen. Ich habe den Raum nie gesehen und keine zwölf Quadratmeter gemessen. Die Zahl war ein Vergleich mit einer früheren Wohnung. Auch meine Aussage, es gehe meiner Frau besser, sollte keine ärztliche Entwarnung sein. Ich habe ihr den beigefügten Brief geschrieben, damit sie weiß, was ich erzählt habe. Sie soll entscheiden, welche persönlichen Details an die Kanzlei gehen. Bitte fragen Sie sie selbst, ob sie eine gemeinsame Unterbringung weiterhin möchte. Ich könnte mir denken, dass ihr Gesellschaft hilft, aber ich bin nicht an ihrer Stelle. Freundliche Grüße, Idris Finkbeiner.''',[14]),
        (27,'2026-10-14','Daria Mertens','abteilung@jva-silberrain.example',c['lawmail'],'Ausgabe von Ersatzmaterial', '''Sehr geehrter Herr Lindenfels, ich übersende die Beobachtungsnotiz zu meiner Schicht. Die erste Schutzdecke war laut Materialliste vorhanden. Eine zusätzliche wurde am Morgen angefordert und ausgegeben. Ob die erste Versorgung ausreichend war, kann ich nicht allein aus dem Wort „ausgegeben“ beurteilen. Eine Temperaturmessung ist nicht eingetragen. Bitte berücksichtigen Sie, dass mein Vermerk „ruhig“ nur die eigene Beobachtung beschreibt und nicht eine ärztliche Entscheidung ersetzt. Die Nachtliste ist eine andere Unterlage. Wir suchen sie heraus und lassen sie über die zuständige Stelle zugänglich machen. Mit freundlichen Grüßen, Daria Mertens.''',[11]),
        (28,'2026-10-15',c['lawyer'],c['lawmail'],'leitung@jva-silberrain.example','Konkrete Unterlagen und aktueller Stand', '''Sehr geehrte Frau Seibold, bitte teilen Sie den aktuellen Stand jeder noch bestehenden besonderen Maßnahme mit. Benötigt werden die Anordnungen und Änderungen, vollständige Beobachtungs- und Materiallisten, die ärztliche Beteiligung sowie die dokumentierten Alternativen zur zeitweisen Dreifachbelegung. Wir unterstellen weder das Fehlen medizinischer Hilfe noch eine uneingeschränkte Rückgabefähigkeit sämtlicher Gegenstände. Bitte sichern Sie die unterschiedlichen Raumplanstände und etwaige Fotos mit ihrer Datierung. Die Mandantin möchte den Schutz ihrer Privatsphäre besprechen. Dieses Anliegen soll nicht als Ablehnung notwendiger Hilfe verstanden werden. Mit freundlichen Grüßen, Xaver Lindenfels.''',[]),
        (29,'2026-10-16','Annemarie Seibold','leitung@jva-silberrain.example',c['lawmail'],'Zwischennachricht zu Ihrer Anfrage', '''Sehr geehrter Herr Lindenfels, die derzeitige Doppelbelegung besteht fort; die dritte Person wurde am 06.10. verlegt. Gewöhnliche Bettwäsche wird seit dem 04.10. verwendet. Die Dokumentation zum Gürtel wird heute mit dem medizinischen Dienst abgeglichen. Wir möchten daraus keine undatierte pauschale Bestätigung machen. Die Reparatur an der Sanitärtür ist beauftragt, aber noch nicht als erledigt zurückgemeldet. Ein Gespräch zur Privatsphäre ist für den 19.10. vorgesehen, sofern Frau Finkbeiner damit einverstanden ist. Die persönliche Anhörung der Mandantin und der Mitgefangenen wird nicht durch diesen Schriftverkehr ersetzt. Mit freundlichen Grüßen, Annemarie Seibold.''',[15]),
        (30,'2026-10-16','Elif Buchner','assistenz@kanzlei-lindenfels.example',c['lawmail'],'Vollmacht und noch offene Rückfragen', '''Herr Lindenfels, die Vollmacht ist eingegangen. Ich habe sie der Akte beigefügt. In unserer Übersicht sind Zustellung an die Voranwältin und spätere Ausgabe an die Mandantin getrennt erfasst. Für die Eingabe vom 13.10. fehlt eine Antwort des Gerichts. Ich habe keine neue Frist aus dem Wort „Beschwerde“ berechnet. Die Mandantin möchte den fertigen Text vorgelesen bekommen; dafür soll ein Gesprächstermin vereinbart werden. Die Tabellen mit den Raumangaben lassen beide Planstände stehen. Bitte markieren Sie nach der fachlichen Prüfung, welche Unterlagen als Anlagen eingereicht werden sollen. Elif Buchner.''',[2]),
        (31,'2026-10-15','Paula Wetter','belegung@jva-silberrain.example',c['lawmail'],'Arbeitsübersicht Belegung und Öffnungszeiten', '''Sehr geehrter Herr Lindenfels, beigefügt ist die tabellarische Arbeitsübersicht aus dem Belegungsdienst. Die Zeiträume sind aus den bezeichneten Übergaben übernommen. Die täglich eingetragenen Zeiten außerhalb des Raums stammen aus dem Stationsjournal; ein leeres Feld bedeutet fehlende Übertragung, nicht null Minuten. Die Übersicht enthält keine medizinische Beurteilung und entscheidet nicht, ob die Fläche rechtlich genügt. Die beiden Messstände bleiben getrennt. Falls Sie einen einzelnen Tag beanstanden, können wir den passenden Originaljournalauszug anfordern. Bitte verwenden Sie die Summen nur unter Beachtung der Erläuterung. Mit freundlichen Grüßen, Paula Wetter.''',[34]),
        (32,'2026-10-16','Sigrun Kappel','s.kappel@kanzlei-kappel.example',c['lawmail'],'Aktenwechsel und chronologische Übersicht', '''Sehr geehrter Kollege Lindenfels, anbei sind mein Übergabevermerk und die von Ihrer Mitarbeiterin zusammengestellte Chronologie mit meinen bestätigten Versanddaten. Die Tatsachen aus der JVA habe ich damit nicht persönlich bestätigt. Meine erste Eingabe enthielt noch keine vollständigen ärztlichen Unterlagen und keinen Raumplan. Ich bitte, die Ausgangsfassungen zu erhalten, damit der Verlauf nachvollziehbar bleibt. Die gesonderte Honorarvereinbarung wird in einem zweiten Übermittlungsweg nachgereicht. Die Mandantin hat mich nicht zu einem Vergleich oder zur Rücknahme ermächtigt. Mit kollegialen Grüßen, Sigrun Kappel.''',[7,33])]
    for n,date,sender,frommail,to,subject,body,attachments in rows:
        ascii_subject=subject.translate(str.maketrans({'ä':'ae','ö':'oe','ü':'ue','Ä':'Ae','Ö':'Oe','Ü':'Ue','ß':'ss'}))
        legacy=f'{n:02d}_{re.sub(r"[^A-Za-z0-9]+", "_", subject).strip("_")}.eml'
        filename=f'{n:02d}_{re.sub(r"[^A-Za-z0-9]+", "_", ascii_subject).strip("_")}.eml'
        c['mails'].append(dict(file=filename,legacy_file=legacy,date=date,sender=sender,frommail=frommail,to=to,subject=subject,body=body,attachments=attachments))

def sheets_data(c):
    if c['land']=='Hessen':
        events=[('2026-09-09','Anordnung 17.30; Beginn getrennte Unterbringung',9,'Anstaltsentscheidung'),('2026-09-09','Bekanntgabe 18.15; keine Empfangsunterschrift',10,'Vollständige Aushändigung laut Vermerk'),('2026-09-10','Frühdienstbericht fertiggestellt',11,'Nicht durchgehende eigene Wahrnehmung'),('2026-09-18','Eingabe an JVA',3,'Antragstellerdarstellung'),('2026-09-22','Hauptsache und vorläufiges Begehren eingereicht',6,'Eingang 16.42 laut Postjournal'),('2026-09-24','Nur vorläufiger Antrag abgelehnt',16,'Hauptsache offen'),('2026-09-25','Beschluss an Voranwalt zugestellt',17,'Elektronisches Empfangsbekenntnis'),('2026-09-28','Beschlusskopie an Gefangenen',17,'Andere Funktion als Kanzleizustellung'),('2026-09-29','Telefonat mit Schwester',14,'Datum später korrigiert'),('2026-10-08','Medizinische Auskunft unvollständig',13,'Befundunterlagen angefordert'),('2026-10-09','Fortdauerentscheidung',15,'Bekanntgabe nicht genau dokumentiert'),('2026-10-15','Anwaltliches Gespräch',4,'Rückbestätigung ausstehend'),('2026-10-16','Aktenübernahme',7,'Kein Verzicht und keine Rücknahme')]
        contacts=[('2026-09-09',0,0,'Anordnung/Stationsübernahme S-0909'),('2026-09-10',10,45,'S-0910; Türkontakt, kein geplanter Sozialtermin'),('2026-09-11',5,60,'S-0911'),('2026-09-12',5,0,'S-0912; Angebot vom Gefangenen abgelehnt'),('2026-09-13',None,None,'S-0913 fehlt im übertragenen Auszug'),('2026-09-14',None,None,'S-0914 fehlt; Gefangener behauptet Wunsch'),('2026-09-17',35,60,'Sozialdiensttermin laut eigener Nachricht'),('2026-09-21',10,60,'S-0921'),('2026-09-24',5,45,'S-0924'),('2026-09-29',20,60,'Telefonat, Datum durch Familie korrigiert'),('2026-10-01',10,60,'S-1001'),('2026-10-04',5,0,'S-1004; keine Begründung übertragen'),('2026-10-06',15,60,'Sozialdienst nennt kurzes Gespräch'),('2026-10-08',5,60,'S-1008'),('2026-10-09',10,60,'S-1009'),('2026-10-12',None,60,'Kontaktminuten nicht eingetragen'),('2026-10-15',43,60,'Anwaltsbesuch 14.05 bis 14.48'),('2026-10-16',None,None,'Laufender Tag; keine vollständige Erfassung')]
        room=[['S 4 Bestandsplan',3.8,2.7,1.55,1,'20_Raumblatt_und_Uebergaben.pdf']]
        periods=[['2026-09-09','2026-10-16',1,'Absonderung; Ende offen; Tabelle nur bis Stichtag']]
        explanation='Der Stationsauszug enthält nur die übertragenen Tage, keine lückenlose Tagesdokumentation. Türkontakte, Telefonat, Sozialdienst und anwaltlicher Besuch sind verschiedenartige Kontakte. Eine Summe ist keine Qualitätsbewertung. Fehlende Minuten bleiben leer. Die laufende Absonderung endet nicht automatisch am Tabellenstichtag.'
    else:
        events=[('2026-09-28','Schutzanordnung 18.40',9,'Verschiedene Maßnahmen getrennt'),('2026-09-29','Schriftliche Aushändigung 09.15',10,'Vorherige Erklärung am Vorabend'),('2026-09-29','Zusätzliche Decke 07.35',10,'Erste Decke laut Liste vorhanden'),('2026-09-30','Eingabe an Anstaltsleitung',3,'Fortgesetzte Behandlung gewünscht'),('2026-10-01','Verlegung nach B 12 um 10.20',15,'Ende besonders gesicherter Unterbringung'),('2026-10-02','Gerichtseingang 15.02',17,'Schriftsatzüberschrift nicht Rechtsbehelfprüfung'),('2026-10-03','Dritte Person ab 18.00',12,'Keine Zustimmung der Mandantin verzeichnet'),('2026-10-04','Gewöhnliche Bettwäsche ausgegeben',15,'Gürtel gesondert'),('2026-10-06','Dritte Person verlässt Raum 09.00',12,'Doppelbelegung besteht fort'),('2026-10-07','Aufmaß Raum B 12',20,'Planstandabweichung'),('2026-10-09','Vorläufiger Antrag abgelehnt',16,'Hauptsache offen'),('2026-10-12','Zustellung an Voranwältin',17,'Elektronisches Empfangsbekenntnis'),('2026-10-13','Sofortige Beschwerde übermittelt',17,'Statthaftigkeit nicht dokumentiert'),('2026-10-14','Medizinische Mitteilung',13,'Vollständige Akte folgt'),('2026-10-15','Anwaltliches Gespräch',4,'Maße aus Erinnerung nicht bestätigt'),('2026-10-16','Übernahme durch neue Kanzlei',7,'Kein neues Gerichtsergebnis')]
        contacts=[('2026-10-01',20,60,'B-1001; Wechsel aus Schutzbereich'),('2026-10-02',40,90,'B-1002'),('2026-10-03',30,60,'B-1003; dritte Person ab18.00'),('2026-10-04',15,45,'B-1004; Sonntagsablauf'),('2026-10-05',None,None,'B-1005 im Auszug nicht enthalten'),('2026-10-06',25,75,'B-1006; dritte Person bis09.00'),('2026-10-07',30,90,'B-1007; technischer Termin nicht eingerechnet'),('2026-10-08',40,90,'B-1008'),('2026-10-09',25,60,'B-1009'),('2026-10-10',15,60,'B-1010'),('2026-10-11',15,45,'B-1011'),('2026-10-12',30,90,'B-1012'),('2026-10-13',40,90,'B-1013'),('2026-10-14',30,90,'B-1014'),('2026-10-15',52,90,'B-1015; Anwalt10.10bis11.02'),('2026-10-16',None,None,'Tag noch nicht abgeschlossen')]
        room=[['B 12 Plan 2018',4.0,3.55,1.8,2,'20_Raumblatt_B12_und_Aufmass.pdf'],['B 12 Aufmaß 07.10.',3.96,3.50,1.92,2,'20_Raumblatt_B12_und_Aufmass.pdf'],['B 12 Aufmaß, drei Personen',3.96,3.50,1.92,3,'20_Raumblatt_B12_und_Aufmass.pdf']]
        periods=[['2026-09-28','2026-10-01',1,'Besonders gesicherter Raum; Uhrzeiten 18.40/10.20'],['2026-10-01','2026-10-03',2,'B 12; dritter Einzug erst03.10.18.00'],['2026-10-03','2026-10-06',3,'B 12; dritte Person bis06.10.09.00'],['2026-10-06','2026-10-16',2,'B 12; aktueller Zustand bisStichtag']]
        explanation='Die Periodentabelle zeigt Kalenderdatendifferenzen, keine auf volle Tage gerundeten Haftzeiten. Uhrzeiten stehen in den Belegtexten. Raumflächen werden technisch berechnet; keine Formel entscheidet über Menschenwürde oder Rechtmäßigkeit. Kontaktminuten sind nicht identisch mit Öffnungsminuten. Fehlende Übertragungen bleiben leer und dürfen nicht als null behandelt werden.'
    filemap={int(d['file'][:2]):d['file'] for d in c['docs']}
    c['events']=[[date,text,filemap[n],status] for date,text,n,status in events]
    c['contacts']=contacts;c['rooms']=room;c['periods']=periods;c['explanation']=explanation
    c['workbook_files']=['33_Chronologie_und_Anlagen.xlsx','34_Unterbringung_und_Kontakte.xlsx']

XLSX_JS = r'''
import fs from 'node:fs/promises';
import path from 'node:path';
import {Workbook, SpreadsheetFile} from '@oai/artifact-tool';
const [datafile, root, qa] = process.argv.slice(2);
const cases=JSON.parse(await fs.readFile(datafile,'utf8'));
const serial=s=>Date.parse(s+'T00:00:00Z')/86400000+25569;
const reports=[];
function sheet(wb,name,title,headers,widths){
 const ws=wb.worksheets.add(name);ws.showGridLines=false;
 const end=String.fromCharCode(64+headers.length);
 ws.getRange(`A1:${end}5`).format.font={name:'Arial',size:11,color:'#172B3A'};
 ws.getRange('A2').values=[[title]];ws.getRange('A2').format.font={name:'Arial',size:15,bold:true,color:'#172B3A'};ws.getRange('A2').format.rowHeight=27;
 ws.getRange(`A5:${end}5`).values=[headers];ws.getRange(`A5:${end}5`).format.fill='#334B62';ws.getRange(`A5:${end}5`).format.font={name:'Arial',size:11,bold:true,color:'#FFFFFF'};
 ws.getRange(`A5:${end}5`).format.rowHeight=31;ws.getRange(`A5:${end}5`).format.wrapText=true;ws.getRange(`A5:${end}5`).format.verticalAlignment='center';
 widths.forEach((width,i)=>ws.getRange(`${String.fromCharCode(65+i)}:${String.fromCharCode(65+i)}`).format.columnWidth=width);
 ws.freezePanes.freezeRows(5);return ws;
}
function body(ws,rows,cols){const end=String.fromCharCode(64+cols);if(rows.length)ws.getRange(`A6:${end}${5+rows.length}`).values=rows;const range=ws.getRange(`A6:${end}${5+rows.length}`);range.format.font={name:'Arial',size:11,color:'#172B3A'};range.format.wrapText=true;range.format.verticalAlignment='center';range.format.rowHeight=42;for(let i=0;i<rows.length;i++)if(i%2===1)ws.getRange(`A${6+i}:${end}${6+i}`).format.fill='#F1F4F6';}
function notes(wb,c,kind){const ws=sheet(wb,'Erläuterung','Hinweise zur Benutzung',['Thema','Erläuterung'],[26,97]);const rows=kind===0?[
 ['Bearbeitungsstand','16.10.2026. Zusammengestellt in der übernehmenden Kanzlei. Originalfassungen bleiben erhalten.'],
 ['Chronologie','Entscheidungsdatum, Bekanntgabe, Kanzleizustellung und Weitergabe sind unterschiedliche Ereignisse. Die Datendifferenz zählt Kalendertage bis zum Stand und ist keine Rechtsbehelfsfrist.'],
 ['Anlagen','A1 bis A12 sind vorläufige Kennungen des bearbeitbaren Schriftsatzes. Eine Zuordnung bestätigt weder die Richtigkeit der Quelle noch eine erfolgreiche Einreichung.'],
 ['Offene Felder','Ein leeres Tatsachenfeld ist unbekannt, nicht null und nicht zugestanden. Rückfragen sind anhand der Originalunterlagen zu beantworten.'],
 ['Weitere Arbeit','Nach Eingang neuer Unterlagen Datum, Herkunft und Bezug ergänzen. Eine frühere Aussage wird nicht still überschrieben. Formeln nur auf tatsächliche Datumswerte anwenden.']
 ]:[['Herkunft','Arbeitsauszug der fiktiven Anstalt aus Raumblatt, Belegungsdienst und Stationsjournal. Es handelt sich nicht um eine vollständige Gesundheits- oder Vollzugsakte.'],['Messwerte','Die Raumwerte werden aus dem jeweils bezeichneten Plan übernommen. Gesamtfläche minus Sanitärfläche ist eine technische Rechengröße. Es gibt keine automatische rechtliche Ampel.'],['Zeiträume',c.explanation],['Fehlende Einträge','Leere Kontakt- oder Öffnungsminuten werden nicht als null summiert. Eine Null steht für einen dokumentierten Nullwert. Ein fehlender Tag ist nicht durch einen Nachbartag zu ersetzen.'],['Nächster Schritt','Fehlende Journalauszüge, aktuelle Anordnungen und konkrete Alternativen anfordern. Die Rechnung ersetzt keine Prüfung von Landesrecht, Schutzbedarf, Privatheit oder Rechtsschutz.']];body(ws,rows,2);ws.getRange('A6:B10').format.rowHeight=68;return{ws,rows:10,cols:2};}
for(const c of cases){
 for(let kind=0;kind<2;kind++){
  const wb=Workbook.create(),views=[],checks=[];
  if(kind===0){
   const ws=sheet(wb,'Chronologie',c.person+' Ereignisse',['Datum','Ereignis','Belegdatei','Einordnung'],[17,49,49,47]);
   ws.getRange('A3').values=[['Stand']];ws.getRange('B3').values=[[serial('2026-10-16')]];ws.getRange('B3').setNumberFormat('yyyy-mm-dd');
   body(ws,c.events.map(r=>[serial(r[0]),...r.slice(1)]),4);ws.getRange(`A6:A${5+c.events.length}`).setNumberFormat('yyyy-mm-dd');views.push({ws,rows:5+c.events.length,cols:4});
   const a=sheet(wb,'Anlagen','Vorläufige Zuordnung zum Entwurf',['Anlage','Datei','Gegenstand','Status'],[14,56,44,40]);
   const nums=c.land==='Hessen'?[9,10,3,6,16,15,25,11,12,13,4,2]:[9,10,11,13,20,12,3,6,16,17,4,2];
   const names=new Map([...c.docs,...c.mails].map(d=>[parseInt(d.file),d.file]));
   body(a,nums.map((n,i)=>['A'+(i+1),names.get(n),n===25?'Sozialdienst nennt Alternative':c.docs.find(d=>parseInt(d.file)===n)?.title||'Nachricht','Zuordnung vor Einreichung prüfen']),4);views.push({ws:a,rows:17,cols:4});
   const t=sheet(wb,'Datumsabstand','Zeit seit dokumentiertem Ereignis',['Datum','Ereignis','Kalendertage bis Stand','Bedeutung'],[17,58,25,54]);
   t.getRange('A3').values=[['Stand']];t.getRange('B3').values=[[serial('2026-10-16')]];t.getRange('B3').setNumberFormat('yyyy-mm-dd');
   body(t,c.events.map(r=>[serial(r[0]),r[1],null,'Keine gesetzliche Fristberechnung']),4);t.getRange(`A6:A${5+c.events.length}`).setNumberFormat('yyyy-mm-dd');
   c.events.forEach((r,i)=>{t.getRange('C'+(6+i)).formulas=[['=$B$3-A'+(6+i)]];checks.push({sheet:'Datumsabstand',cell:'C'+(6+i),expected:serial('2026-10-16')-serial(r[0])});});views.push({ws:t,rows:5+c.events.length,cols:4});
  }else{
   const r=sheet(wb,'Raumwerte','Technische Flächenberechnung',['Planstand','Gesamt m²','Sanitär m²','Personen','Restfläche m²','Rest je Person m²'],[38,19,19,15,22,24]);
   body(r,c.rooms.map(v=>[v[0],v[1]*v[2],v[3],v[4],null,null]),6);
   c.rooms.forEach((v,i)=>{let row=6+i;r.getRange('E'+row).formulas=[['=B'+row+'-C'+row]];r.getRange('F'+row).formulas=[['=IF(D'+row+'>0,E'+row+'/D'+row+',"Personenzahl fehlt")']];checks.push({sheet:'Raumwerte',cell:'E'+row,expected:v[1]*v[2]-v[3]},{sheet:'Raumwerte',cell:'F'+row,expected:(v[1]*v[2]-v[3])/v[4]});});r.getRange('B6:C'+(5+c.rooms.length)).setNumberFormat('0.00');r.getRange('E6:F'+(5+c.rooms.length)).setNumberFormat('0.00');
   const sourceRow=8+c.rooms.length;r.getRange(`A${sourceRow}`).values=[['Quelle']];r.getRange(`B${sourceRow}:F${sourceRow}`).merge();r.getRange(`B${sourceRow}`).values=[[c.rooms[0][5]]];r.getRange(`A${sourceRow}:F${sourceRow}`).format.rowHeight=47;views.push({ws:r,rows:sourceRow,cols:6});
   const p=sheet(wb,'Zeiträume','Kalenderdaten der Unterbringung',['Von','Bis Stand oder Wechsel','Personen','Dokumentierter Ablauf'],[18,25,16,91]);body(p,c.periods.map(v=>[serial(v[0]),serial(v[1]),v[2],v[3]]),4);p.getRange('A6:B'+(5+c.periods.length)).setNumberFormat('yyyy-mm-dd');views.push({ws:p,rows:5+c.periods.length,cols:4});
   const k=sheet(wb,'Kontakte','Übertragene Kontakt und Öffnungsminuten',['Datum','Kontaktminuten','Minuten außerhalb','Beleg oder Lücke'],[18,23,24,89]);
   body(k,c.contacts.map(v=>[serial(v[0]),v[1],v[2],v[3]]),4);k.getRange('A6:A'+(5+c.contacts.length)).setNumberFormat('yyyy-mm-dd');
   const sum=7+c.contacts.length;k.getRange('A'+sum).values=[['Erfasste Minuten']];k.getRange('B'+sum).formulas=[['=SUM(B6:B'+(5+c.contacts.length)+')']];k.getRange('C'+sum).formulas=[['=SUM(C6:C'+(5+c.contacts.length)+')']];k.getRange('D'+sum).values=[['Nur vorhandene Einträge; fehlende Tage nicht aufgefüllt']];k.getRange(`A${sum}:D${sum+1}`).format.rowHeight=52;k.getRange(`A${sum}:D${sum+1}`).format.font={name:'Arial',size:11,color:'#172B3A'};k.getRange(`A${sum}:D${sum+1}`).format.wrapText=true;k.getRange(`A${sum}:D${sum+1}`).format.verticalAlignment='center';
   k.getRange('A'+(sum+1)).values=[['Fehlende Felder']];k.getRange('B'+(sum+1)).formulas=[['=COUNTBLANK(B6:B'+(5+c.contacts.length)+')']];k.getRange('C'+(sum+1)).formulas=[['=COUNTBLANK(C6:C'+(5+c.contacts.length)+')']];k.getRange('D'+(sum+1)).values=[['Leere Felder sind unbekannt und keine Nullwerte']];
   checks.push({sheet:'Kontakte',cell:'B'+sum,expected:c.contacts.reduce((s,v)=>s+(v[1]??0),0)},{sheet:'Kontakte',cell:'C'+sum,expected:c.contacts.reduce((s,v)=>s+(v[2]??0),0)},{sheet:'Kontakte',cell:'B'+(sum+1),expected:c.contacts.filter(v=>v[1]===null).length},{sheet:'Kontakte',cell:'C'+(sum+1),expected:c.contacts.filter(v=>v[2]===null).length});views.push({ws:k,rows:sum+1,cols:4});
  }
  views.push(notes(wb,c,kind));wb.recalculate();
  checks.forEach(x=>{x.actual=wb.worksheets.getItem(x.sheet).getRange(x.cell).values[0][0];if(typeof x.actual!=='number'||Math.abs(x.actual-x.expected)>1e-8)throw Error(JSON.stringify(x));});
  // Rechenprobe mit geändertem Eingang, danach ursprünglichen Stand wiederherstellen.
  const target=kind===0?wb.worksheets.getItem('Datumsabstand'):wb.worksheets.getItem('Raumwerte');const testCell=kind===0?'B3':'D6';const resultCell=kind===0?'C6':'F6';const old=target.getRange(testCell).values[0][0];const initial=target.getRange(resultCell).values[0][0];target.getRange(testCell).values=[[old+1]];wb.recalculate();const changed=target.getRange(resultCell).values[0][0];if(changed===initial)throw Error('Keine Neuberechnung');target.getRange(testCell).values=[[old]];wb.recalculate();if(Math.abs(target.getRange(resultCell).values[0][0]-initial)>1e-8)throw Error('Wiederherstellung fehlt');
  const errors=await wb.inspect({kind:'match',searchTerm:'#REF!|#DIV/0!|#VALUE!|#NAME\\?|#N/A|#NUM!|#NULL!|#SPILL!|#CALC!',options:{useRegex:true,maxResults:30},maxChars:2000});
  const file=c.workbook_files[kind],folder=path.join(root,'testakten',c.slug),renderdir=path.join(qa,c.slug,file.replace('.xlsx',''));await fs.mkdir(renderdir,{recursive:true});
  for(const v of views){const blob=await wb.render({sheetName:v.ws.name,range:`A1:${String.fromCharCode(64+v.cols)}${v.rows}`,scale:1,format:'png'});await fs.writeFile(path.join(renderdir,v.ws.name+'.png'),new Uint8Array(await blob.arrayBuffer()));}
  const output=await SpreadsheetFile.exportXlsx(wb);await output.save(path.join(folder,file));await fs.rm(path.join(folder,file+'.inspect.ndjson'),{force:true});
  reports.push({slug:c.slug,file,sheets:views.map(v=>({name:v.ws.name,range:`A1:${String.fromCharCode(64+v.cols)}${v.rows}`})),formula_checks:checks,input_change_test:{cell:testCell,before:initial,after:changed,restored:true},errors:errors.ndjson});
 }
}
await fs.writeFile(path.join(qa,'xlsx-pruefung.json'),JSON.stringify(reports,null,2)+'\n');console.log('Vier Arbeitsmappen erstellt, alle Blätter gerendert, Formeln und Eingabeänderungen geprüft.');
'''

def write_documents(c):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak
    from reportlab.graphics.shapes import Drawing, Rect, Line, String
    folder=ROOT/'testakten'/c['slug'];folder.mkdir(parents=True,exist_ok=True)
    fonts=Path('/System/Library/Fonts/Supplemental')
    if not (fonts/'Times New Roman.ttf').exists():
        fonts=Path('/usr/share/fonts/truetype/liberation2');names=['LiberationSerif-Regular.ttf','LiberationSerif-Bold.ttf']
    else:names=['Times New Roman.ttf','Times New Roman Bold.ttf']
    for key,name in zip(['JvaErw','JvaErwBold'],names):pdfmetrics.registerFont(TTFont(key,str(fonts/name)))
    ps=ParagraphStyle('body',fontName='JvaErw',fontSize=11,leading=14.1,spaceAfter=9)
    hs=ParagraphStyle('heading',parent=ps,fontName='JvaErwBold',spaceBefore=7,spaceAfter=10,keepWithNext=True)
    ts=ParagraphStyle('title',parent=hs,fontSize=16,leading=19,spaceAfter=17)
    for item in c['docs']:
        target=folder/item['file']
        if target.suffix=='.docx':
            d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(22);s.left_margin=s.right_margin=Mm(23)
            for name in ['Normal','Title','Heading 1']:
                st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
                for attr in list(st.element.rPr.rFonts.attrib):
                    if attr.endswith('Theme') or attr.endswith('theme'):del st.element.rPr.rFonts.attrib[attr]
                for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
            d.styles['Normal'].paragraph_format.space_after=Pt(7 if item['file'].startswith('08_') else 9);d.styles['Normal'].paragraph_format.line_spacing=1.04 if item['file'].startswith('08_') else 1.08
            d.styles['Title'].font.size=Pt(16);d.styles['Title'].font.bold=True
            d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(10);d.styles['Heading 1'].paragraph_format.space_after=Pt(10)
            d.add_paragraph(item['title'],'Title')
            for p in item['paragraphs']:d.add_paragraph(p,'Heading 1' if re.match(r'^\d+\. ',p) and len(p)<100 else 'Normal')
            footer=s.footer.paragraphs[0];footer.alignment=2;run=footer.add_run(c['record']+'  |  ');run.font.name='Times New Roman';run.font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fieldrun=OxmlElement('w:r');fieldtext=OxmlElement('w:t');fieldtext.text='1';fieldrun.append(fieldtext);field.append(fieldrun);footer._p.append(field)
            d.core_properties.author=c['prison'];d.core_properties.title=item['title'];d.core_properties.language='de-DE';d.core_properties.created=dt.datetime(2026,10,16);d.core_properties.modified=dt.datetime(2026,10,16);d.save(target)
        else:
            flow=[Paragraph(escape(item['title']),ts)]
            for p in item['paragraphs']:flow.append(Paragraph(escape(p).replace('\n','<br/>'),hs if re.match(r'^\d+\. ',p) and len(p)<100 else ps))
            if c['land']=='Baden-Württemberg' and item['file'].startswith('20_'):
                flow.extend([PageBreak(),Paragraph('Raum B 12 nach Entfernung des dritten Bettes',hs),Paragraph('Schematische Orientierung zum Aufmaß vom 07.10.2026. Nicht maßstabsgerecht. Die Zeichnung zeigt den späteren Zustand und belegt keine frühere Möblierung.',ps)])
                diagram=Drawing(450,290)
                diagram.add(Rect(50,40,350,210,strokeWidth=2,fillColor=None))
                diagram.add(Rect(50,170,95,80,strokeWidth=1,fillColor=None));diagram.add(String(65,208,'Sanitärzone',fontName='JvaErw',fontSize=11));diagram.add(String(65,190,'1,92 m²',fontName='JvaErw',fontSize=11))
                diagram.add(Rect(60,52,70,100,strokeWidth=1,fillColor=None));diagram.add(String(68,99,'Bett 1',fontName='JvaErw',fontSize=11))
                diagram.add(Rect(318,138,70,100,strokeWidth=1,fillColor=None));diagram.add(String(326,187,'Bett 2',fontName='JvaErw',fontSize=11))
                diagram.add(Rect(195,128,75,52,strokeWidth=1,fillColor=None));diagram.add(String(216,150,'Tisch',fontName='JvaErw',fontSize=11))
                diagram.add(Line(185,40,240,40,strokeWidth=5,strokeColor=__import__('reportlab.lib.colors',fromlist=['white']).white));diagram.add(String(188,18,'Tür',fontName='JvaErw',fontSize=11))
                diagram.add(String(200,269,'Innenlänge 3,96 m',fontName='JvaErw',fontSize=11));diagram.add(String(55,5,'Innenbreite 3,50 m; Bestandsplan und Aufmaß nicht vermischen.',fontName='JvaErw',fontSize=10))
                flow.append(diagram)
                flow.append(Paragraph('Die Position des vorübergehenden dritten Bettes ist in dieser Skizze nicht rekonstruiert. Fotos oder eine datierte Belegungsskizze für den 03. bis 06.10. sind angefordert. Die Sanitärtür wurde bei der Besichtigung als nicht vollständig schließend beschrieben.',ps))
            def footer(can,document):can.setFont('JvaErw',9);can.drawString(54,27,c['record']);can.drawRightString(A4[0]-54,27,str(document.page))
            SimpleDocTemplate(str(target),pagesize=A4,leftMargin=54,rightMargin=54,topMargin=48,bottomMargin=48,invariant=1,title=item['title'],author=c['prison']).build(flow,onFirstPage=footer,onLaterPages=footer)

def write_mails(c):
    folder=ROOT/'testakten'/c['slug'];files={int(d['file'][:2]):d['file'] for d in c['docs']};files.update({int(f[:2]):f for f in c['workbook_files']})
    for item in c['mails']:
        if item['legacy_file']!=item['file']:
            old=folder/item['legacy_file']
            if old.exists():old.unlink()
        msg=EmailMessage(policy=SMTP);msg['From']=f'{item["sender"]} <{item["frommail"]}>';msg['To']=item['to'];msg['Subject']=item['subject'];msg['Date']=format_datetime(dt.datetime.fromisoformat(item['date']+'T10:30:00').replace(tzinfo=ZoneInfo('Europe/Berlin')));msg['Message-ID']=f'<{c["slug"]}-{item["file"][:2]}@{c["domain"]}>';msg['Content-Language']='de-DE';msg.set_content(item['body'])
        for number in item['attachments']:
            name=files[number];p=folder/name
            subtype={'.pdf':'pdf','.docx':'vnd.openxmlformats-officedocument.wordprocessingml.document','.xlsx':'vnd.openxmlformats-officedocument.spreadsheetml.sheet'}[p.suffix]
            msg.add_attachment(p.read_bytes(),maintype='application',subtype=subtype,filename=name)
        if item['attachments']:msg.set_boundary('JVA-ERW-'+c['slug']+'-'+item['file'][:2])
        (folder/item['file']).write_bytes(msg.as_bytes())

def print_profiles(cases):
    """Druckmetadaten ergänzen; Werte, Formeln und Formatierung nicht verändern."""
    import io, zipfile, xml.etree.ElementTree as ET
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';ET.register_namespace('',ns)
    reports=json.loads((QA/'xlsx-pruefung.json').read_text())
    for c in cases:
        for filename in c['workbook_files']:
            p=ROOT/'testakten'/c['slug']/filename
            config=next(r for r in reports if r['slug']==c['slug'] and r['file']==filename)
            out=io.BytesIO()
            with zipfile.ZipFile(p) as src,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
                for item in src.infolist():
                    data=src.read(item)
                    if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',item.filename):
                        tree=ET.fromstring(data);setup=tree.find('{'+ns+'}pageSetup')
                        if setup is None:setup=ET.SubElement(tree,'{'+ns+'}pageSetup')
                        setup.set('orientation','landscape');setup.set('paperSize','9');setup.set('fitToWidth','1');setup.set('fitToHeight','0')
                        margin=tree.find('{'+ns+'}pageMargins')
                        if margin is None:margin=ET.SubElement(tree,'{'+ns+'}pageMargins')
                        for edge in ['left','right','top','bottom']:margin.set(edge,'0.3')
                        margin.set('header','0.15');margin.set('footer','0.15')
                        data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
                    elif item.filename=='xl/workbook.xml':
                        tree=ET.fromstring(data);defined=tree.find('{'+ns+'}definedNames')
                        if defined is None:
                            defined=ET.Element('{'+ns+'}definedNames');calc=tree.find('{'+ns+'}calcPr');tree.insert(list(tree).index(calc) if calc is not None else len(tree),defined)
                        for i,v in enumerate(config['sheets']):
                            for old in list(defined):
                                if old.get('localSheetId')==str(i) and old.get('name') in ['_xlnm.Print_Area','_xlnm.Print_Titles']:defined.remove(old)
                            area=ET.SubElement(defined,'{'+ns+'}definedName',name='_xlnm.Print_Area',localSheetId=str(i));area.text="'"+v['name']+"'!"+v['range']
                            repeated=ET.SubElement(defined,'{'+ns+'}definedName',name='_xlnm.Print_Titles',localSheetId=str(i));repeated.text="'"+v['name']+"'!$5:$5"
                        data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
                    dst.writestr(item,data)
            p.write_bytes(out.getvalue())

def write_readme(c):
    text=f'''# Testakte {c['title']}

[Alle Testakten](../README.md) · [JVA-Plugin](../../strafvollzug-jva/README.md) · [Hauptübersicht](../../README.md) · [Download-Index](../../ASSET_INDEX.md)

Diese Arbeitsakte lässt sich aus Sicht der betroffenen Person, ihrer anwaltlichen Vertretung, der JVA oder der Strafvollstreckungskammer bearbeiten. Beginnen Sie mit dem Mandatsauftrag und rekonstruieren Sie anschließend die einzelnen Maßnahmen und den gerichtlichen Verlauf. Stand ist der 16.10.2026. Die ursprünglichen Anträge bleiben unverändert erhalten; der letzte Word-Entwurf ist zur weiteren Bearbeitung bestimmt und noch nicht versandfertig.

## 1. Inhalt und Benutzung

Die Akte enthält **34 Originaldateien**: acht Word-Dokumente, zwölf PDFs, zwölf E-Mails mit echten PDF-, Word- und Excel-Anhängen sowie zwei bearbeitbare Excel-Arbeitsmappen. Die Anhänge sind zusätzliche Kopien vorhandener Aktenstücke und werden nicht als weitere Originale gezählt. Die Arbeitsmappen enthalten zusammen acht Blätter mit Erklärungen, Chronologie, vorläufigen Anlagenkennungen, Raumwerten, Zeitabschnitten und Kontaktangaben. Ihre Formeln rechnen Daten und Maße; sie entscheiden keine Rechtsfrage.

Ermitteln Sie zunächst, wer was wann angeordnet, beobachtet, weitergegeben und beantragt hat. Unterscheiden Sie Schutzbedarf und dessen Durchführung, Tatsachenbehauptung und medizinischen Befund sowie Hauptsache und vorläufigen Rechtsschutz. Prüfen Sie die ursprünglichen anwaltlichen Eingaben und den gerichtlichen Beschluss anhand der später hinzugekommenen Unterlagen. Ergänzen Sie anschließend die ausdrücklich offenen Tatsachen- und Antragsfelder des letzten Entwurfs. Die Akte enthält keine Musterlösung.

## 2. Landesrecht und Quellen

Die Vorgänge betreffen rechtskräftige Erwachsenenstrafhaft in {c['land']}. Materielle Landesnormen und bundesrechtlicher gerichtlicher Rechtsschutz sind getrennt zu bestimmen. Hausordnungen, Dienstanweisungen, Anstaltsbescheide und die erfundenen gerichtlichen Beschlüsse sind Sachverhaltsbelege, keine amtlichen Rechtsquellen. Die gesetzliche und rechtsprechungsbezogene Ausgangsbasis steht im bearbeitbaren Schriftsatz. Die amtlichen Recherchezugänge finden Sie zusätzlich im [Landesrecht-Wegweiser](../../strafvollzug-jva/references/landesrecht-wegweiser.md) und bei den [Haftbedingungen und Schutzmaßnahmen](../../strafvollzug-jva/references/haftbedingungen-und-schutzmassnahmen.md). Maßnahmetag und aktuelle Fassung müssen zusammenpassen.

## 3. Fiktive Personen und sensible Inhalte

Personen, Anstalt, Anschriften, Aktenzeichen, Nachrichten und medizinische Vorgänge sind erfunden. Gerichtliche Aktenzeichen tragen ausdrücklich den Zusatz „ÜBUNG“. Die Vorwürfe gegen Bedienstete sind unbewiesene Falldarstellungen; es werden keine Vorwürfe gegen wirkliche Anstalten oder Personen erhoben. Die Akte beschreibt Schutzbedarf ohne Selbstschädigungsmethoden. Gefangenen wird kein freier E-Mail-Zugang unterstellt. Die reservierten `.example`-Adressen sind keine Versandziele.

<!-- reserved-example-contacts -->

<!-- BEGIN gesamt-pdf-section (autogen) -->
## 4. Downloads

> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.
>
> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.

| Format | Download |
| --- | --- |
| Gesamt-PDF | [Lesefassung](gesamt-pdf/{c['slug']}_gesamt.pdf) |
| Akten-ZIP | [Bearbeitbare Originaldateien](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/strafvollzug-jva-v1.1.0/testakte-{c['slug']}.zip) |
| Einzel-PDF-ZIP | [Alle Dokumente einzeln als PDF](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/strafvollzug-jva-v1.1.0/testakte-{c['slug']}-einzelpdfs.zip) |

<!-- END gesamt-pdf-section (autogen) -->
'''
    folder=ROOT/'testakten'/c['slug'];(folder/'README.md').write_text(text,encoding='utf-8')
    (folder/'rubric.yaml').write_text(f'''name: {c['slug']}
plugin: strafvollzug-jva
description: "{c['title']}: Alltag, Unterlagen und fortzuführendes gerichtliches Verfahren."
checks:
  - id: originale
    check_type: working_file_count
    min: 34
    description: "Native Arbeitsakte mit Schriftverkehr, Anlagen und Tabellen."
  - id: zeitraeume
    check_type: human_review
    description: "Die einzelnen Maßnahmen, Zeitabschnitte und Erkenntnisquellen werden getrennt bearbeitet."
  - id: verfahren
    check_type: human_review
    description: "Die ursprünglichen Eingaben und der gerichtliche Stand werden ohne automatische Rechtsbehelfsannahme geprüft."
  - id: landesrecht
    check_type: human_review
    description: "Vollzugsart, Landesrecht und Bundesrechtsschutz werden anhand amtlicher Quellen bestimmt."
''',encoding='utf-8')

def manifest(cases):
    from pypdf import PdfReader
    from email import policy
    from email.parser import BytesParser
    xlsx=json.loads((QA/'xlsx-pruefung.json').read_text())
    report={'schema_version':1,'stand':'2026-10-16','plugin':'strafvollzug-jva','component_version':'1.1.0','cases':[],'limits':['Rechenprüfung und redaktionelle Inhaltsprüfung sind kein Live-Modelltest.','Original-PDF-Seiten werden geprüft; native Office- und Gesamt-PDF-Seiten prüft zusätzlich der Releasebuilder.']}
    for c in cases:
        folder=ROOT/'testakten'/c['slug'];files=[];mails=[]
        for p in sorted(folder.iterdir()):
            if p.suffix not in ('.pdf','.docx','.xlsx','.eml'):continue
            entry={'file':p.name,'type':p.suffix[1:],'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size}
            if p.suffix=='.pdf':entry['pages']=len(PdfReader(p).pages)
            files.append(entry)
            if p.suffix=='.eml':
                msg=BytesParser(policy=policy.default).parsebytes(p.read_bytes());attachments=[]
                for a in msg.iter_attachments():
                    name=a.get_filename();data=a.get_payload(decode=True);assert data==(folder/name).read_bytes(),name
                    attachments.append({'filename':name,'sha256':hashlib.sha256(data).hexdigest()})
                mails.append({'file':p.name,'attachments':attachments})
        assert len(files)==34,(c['slug'],len(files))
        report['cases'].append({'slug':c['slug'],'original_count':len(files),'counts':{ext:sum(f['type']==ext for f in files) for ext in ['docx','pdf','eml','xlsx']},'files':files,'mails':mails,'workbooks':[r for r in xlsx if r['slug']==c['slug']],'draft':'08_'+c['docs'][7]['file'][3:],'attachment_map':{'A'+str(i+1):next(d['file'] for d in c['docs']+c['mails'] if int(d['file'][:2])==n) for i,n in enumerate([9,10,3,6,16,15,25,11,12,13,4,2] if c['land']=='Hessen' else [9,10,11,13,20,12,3,6,16,17,4,2])}})
    layout_path=QA/'layout-review.json'
    if layout_path.exists():
        layout=json.loads(layout_path.read_text())
        current={f"testakten/{case['slug']}/{item['file']}":item['sha256'] for case in report['cases'] for item in case['files']}
        if all(current.get(item['source'])==item['source_sha256'] for item in layout['files']):
            report['layout_review']=layout
    out=ROOT/'quality/strafvollzug-jva/erweiterungsakten-pruefung.json';out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(out)

def main():
    p=argparse.ArgumentParser();p.add_argument('--documents-only',action='store_true');p.add_argument('--manifest-only',action='store_true');p.add_argument('--readmes-only',action='store_true');args=p.parse_args()
    QA.mkdir(parents=True,exist_ok=True)
    cases=[hessen(),bw()]
    for c in cases:
        populate_mails(c);sheets_data(c)
        if c['land']=='Hessen':
            c['mails'][0]['date']='2026-09-23'
            for m in c['mails']:
                if int(m['file'][:2]) in [23,24,25,26,27]:m['date']='2026-10-16'
            c['docs'][0]['paragraphs']=[p.replace('Frau Rehbein übernimmt heute die Vertretung vom Kollegen Volker Brack.','Frau Rehbein hat die Vertretung am 15.10. übernommen; die vollständige Übergabe des Kollegen Volker Brack erfolgt heute.') for p in c['docs'][0]['paragraphs']]
            ps=c['docs'][7]['paragraphs']
            insert=next(i for i,p in enumerate(ps) if p.startswith('Der verfassungsrechtliche Maßstab'))
            ps.insert(insert,'§ 50 Absatz 7 HStVollzG knüpft eine Absonderung über 24 Stunden an ihre Unerlässlichkeit. Absatz 8 regelt besondere Betreuung sowie bei mehr als 30 Tagen Dauer oder insgesamt mehr als 90 Tagen innerhalb von zwölf Monaten die Zustimmung der Aufsichtsbehörde. § 51 Absatz 2 verlangt bei einer Absonderung über 24 Stunden spätestens nach drei Tagen eine ärztliche Stellungnahme und danach angemessene Abstände; Absatz 3 verlangt fortlaufende Überprüfung, Absatz 5 Dokumentation. Die amtliche Gesamtausgabe wurde am 16.10.2026 geprüft. Ob die konkreten Anforderungen eingehalten wurden, ist anhand der vollständigen Unterlagen zu bestimmen. Eine lediglich abgesandte Anfrage belegt noch keine erteilte Zustimmung.')
            insert=next(i for i,p in enumerate(ps) if p.startswith('Der Hinweis auf fehlende Unterlagen'))
            ps.insert(insert,'Der Beschluss des Bundesverfassungsgerichts vom 15.01.2020, Az. 2 BvR 1763/16, Rn. 39–43, betrifft effektive Strafverfolgung im Zusammenhang mit einer psychiatrischen Fixierung. Er berücksichtigt auch die besondere Beweissituation in staatlichem Gewahrsam. Er stellt keine Misshandlung in diesem fiktiven Fall fest und ersetzt keine Prüfung des richtigen Verfahrens. Ein etwaiger gesonderter Ermittlungsauftrag ist mit dem Mandanten noch zu klären. Die streitigen Wahrnehmungen und die medizinische Dokumentationslücke sollen deshalb gesichert, nicht als fertiger Schuldnachweis ausgegeben werden.')
            ps.extend(['6. Amtliche Recherchezugänge','HStVollzG: https://www.rv.hessenrecht.hessen.de/bshe/document/jlr-StVollzGHEV6IVZ. BVerfG, Beschluss vom 15.01.2020, Az. 2 BvR 1763/16: https://www.bundesverfassungsgericht.de/SharedDocs/Entscheidungen/DE/2020/01/rk20200115_2bvr176316.html. BVerfG, Beschluss vom 20.04.2026, Az. 2 BvR 1650/24: https://www.bundesverfassungsgericht.de/SharedDocs/Entscheidungen/DE/2026/04/rk20260420_2bvr165024.html. Die Randnummern und die Übertragbarkeit sind bei Fertigstellung erneut am amtlichen Volltext abzugleichen.'])
        else:
            for m in c['mails']:
                if int(m['file'][:2])==24:m['date']='2026-10-16'
            ps=c['docs'][7]['paragraphs'];insert=next(i for i,p in enumerate(ps) if p.startswith('Das Bundesverfassungsgericht verlangt'))
            ps.insert(insert,'§ 67 Absatz 5 JVollzGB Buch 3 begrenzt besondere Sicherungsmaßnahmen auf die erforderliche Dauer. Nach § 70 Absatz 2 ist bei psychischem Zustand als Anlass vor der Anordnung ärztlicher Rat einzuholen; bei Gefahr im Verzug ist die Beteiligung unverzüglich nachzuholen. Für den besonders gesicherten Haftraum verlangt § 71 Absatz 1 alsbaldigen und anschließend möglichst täglichen ärztlichen Besuch. § 68 enthält besondere Voraussetzungen für die dort geregelte Einzelhaft. Diese Anforderungen sind voneinander und von der Nachtunterbringung nach § 13 zu unterscheiden. Die am 16.10.2026 amtlich gelesene Gesamtfassung berücksichtigt die Änderung vom 10.02.2026 (GBl. Nr. 21).')
            insert=next(i for i,p in enumerate(ps) if p.startswith('Die ärztliche Notiz beschreibt'))
            ps.insert(insert,'Der Beschluss des Bundesverfassungsgerichts vom 18.03.2015, Az. 2 BvR 1111/13, Rn. 29–36, verlangt eine gesonderte Rechtfertigung einzelner Sicherungsmaßnahmen und befasst sich mit konkret erforderlicher Entkleidung und Ersatzkleidung. Daraus folgt keine pauschale Freigabe gewöhnlicher Bettwäsche in jedem Gefährdungsfall. Der Beschluss vom 24.01.2008, Az. 2 BvR 1661/06, Rn. 38–43, behandelt tatsächliche Gefahr, fachärztliche Befassung und weniger belastende Alternativen im damaligen Untersuchungshaftkontext. Er kann nur mit offengelegter Übertragungsgrenze als verfassungsrechtliche Orientierung dienen; das heutige baden-württembergische Landesrecht bleibt gesondert anzuwenden.')
            ps.extend(['6. Amtliche Recherchezugänge','JVollzGB Buch 3: https://www.landesrecht-bw.de/bsbw/document/jlr-NNLBW00007CFA. BVerfG, Beschluss vom 18.03.2015, Az. 2 BvR 1111/13: https://www.bundesverfassungsgericht.de/SharedDocs/Entscheidungen/DE/2015/03/rk20150318_2bvr111113.html. BVerfG, Beschluss vom 24.01.2008, Az. 2 BvR 1661/06: https://www.bundesverfassungsgericht.de/SharedDocs/Entscheidungen/DE/2008/01/rk20080124_2bvr166106.html. Weitere zitierte Entscheidungen sind über die amtliche Entscheidungssuche mit Datum und Aktenzeichen aufzurufen. Die Randnummern und der Verfahrenskontext sind vor Einreichung abzugleichen.'])
    if args.manifest_only:manifest(cases);return
    if args.readmes_only:
        for c in cases:write_readme(c)
        return
    for c in cases:write_documents(c);write_readme(c)
    if args.documents_only:return
    (QA/'data.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2)+'\n')
    (QA/'workbooks.mjs').write_text(XLSX_JS)
    modules=QA/'node_modules'
    if not modules.exists():modules.symlink_to(RUNTIME/'node/node_modules',target_is_directory=True)
    subprocess.run([str(RUNTIME/'node/bin/node'),str(QA/'workbooks.mjs'),str(QA/'data.json'),str(ROOT),str(QA)],check=True)
    print_profiles(cases)
    for c in cases:write_mails(c)
    manifest(cases)

if __name__=='__main__':main()
