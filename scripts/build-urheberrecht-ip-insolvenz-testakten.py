#!/usr/bin/env python3
"""Erzeugt drei offene IP-Insolvenz-Testakten aus strukturierten Falldaten.

Die Akten enden jeweils beim unterschriebenen Asset Deal und der angekündigten
Verwertung. Eine anwaltliche Lösung, ein Gutachten oder eine Anspruchsstrategie
werden bewusst nicht erzeugt. README und Rubrik bleiben redaktionelle Dateien
und werden von den zentralen Aktenexporten ausgeschlossen.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import mimetypes
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from email.message import EmailMessage
from email.parser import BytesParser
from email.policy import SMTP
from email.utils import format_datetime, parsedate_to_datetime
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from testakte_disclaimer import NOTICE_MARKDOWN


ROOT = Path(__file__).resolve().parent.parent
TESTAKTEN = ROOT / "testakten"
DATA = ROOT / "scripts" / "data"
MODULES = (
    DATA / "urheberrecht_ip_insolvenz_marken.py",
    DATA / "urheberrecht_ip_insolvenz_software.py",
    DATA / "urheberrecht_ip_insolvenz_film.py",
)
RELEASE_TAG = "urheberrecht-ip-insolvenz-testakten-v1.0.1"
RELEASE_BASE = (
    "https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/"
    + RELEASE_TAG
)
ARTIFACT_MANIFEST = ROOT / "quality" / "legal" / "urheberrecht-ip-insolvenz-artifact-manifest.json"


def load_case(path: Path) -> dict[str, Any]:
    spec = importlib.util.spec_from_file_location(path.stem, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Falldaten nicht importierbar: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    case = module.CASE
    if not isinstance(case, dict):
        raise TypeError(f"{path}: CASE ist kein Dictionary")
    required = {"slug", "title", "short_title", "plugin", "summary", "scenario", "parties", "questions", "ambiguities", "documents"}
    missing = required - set(case)
    if missing:
        raise ValueError(f"{path}: CASE-Felder fehlen: {sorted(missing)}")
    return case


def cases() -> list[dict[str, Any]]:
    missing = [path for path in MODULES if not path.is_file()]
    if missing:
        raise SystemExit("Falldaten fehlen: " + ", ".join(str(path) for path in missing))
    loaded = [load_case(path) for path in MODULES]
    slugs = [case["slug"] for case in loaded]
    if len(set(slugs)) != len(slugs):
        raise ValueError("Doppelte Akten-Slugs")
    return loaded


def safe_rel(value: str) -> Path:
    rel = Path(value)
    if rel.is_absolute() or not rel.parts or ".." in rel.parts or any(part.startswith(".") for part in rel.parts):
        raise ValueError(f"Unzulässiger Aktenpfad: {value!r}")
    return rel


def clear_case(case: dict[str, Any]) -> Path:
    folder = TESTAKTEN / case["slug"]
    if folder.exists():
        shutil.rmtree(folder)
    folder.mkdir(parents=True)
    return folder


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_doc(document: Document, title: str, subject: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.15)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    for name, size, color in (
        ("Title", 18, "183A55"),
        ("Heading 1", 14, "183A55"),
        ("Heading 2", 12, "3C5365"),
        ("Heading 3", 11, "3C5365"),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
    props = document.core_properties
    props.title = title
    props.subject = subject
    props.author = "Fiktive Aktenstelle"
    props.keywords = "Urheberrecht, Immaterialgüterrecht, Insolvenzrecht"
    props.created = datetime(2026, 10, 23, 9, 0, 0)
    props.modified = datetime(2026, 10, 23, 9, 0, 0)
    header = section.header.paragraphs[0]
    header.text = subject
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 96, 100)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Seite ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    for run in footer.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)


def add_paragraph(document: Document, text: str, *, bold_lead: bool = False) -> None:
    paragraph = document.add_paragraph()
    if bold_lead and ":" in text:
        lead, rest = text.split(":", 1)
        run = paragraph.add_run(lead + ":")
        run.bold = True
        paragraph.add_run(rest)
    else:
        paragraph.add_run(text)


def add_table(document: Document, table_spec: dict[str, Any]) -> None:
    headers = [str(value) for value in table_spec.get("headers", [])]
    rows = table_spec.get("rows", [])
    if not headers:
        return
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, "DCE6ED")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
    set_repeat_table_header(table.rows[0])
    for raw_row in rows:
        values = list(raw_row)
        row = table.add_row()
        for index, cell in enumerate(row.cells):
            value = values[index] if index < len(values) else ""
            cell.text = "" if value is None else str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
    document.add_paragraph()


def build_docx(path: Path, case: dict[str, Any], spec: dict[str, Any]) -> None:
    document = Document()
    configure_doc(document, spec["title"], case["short_title"])
    document.add_heading(spec["title"], 0)
    for paragraph in spec.get("paragraphs", []):
        add_paragraph(document, str(paragraph), bold_lead=True)
    for table in spec.get("tables", []):
        add_table(document, table)
    for section in spec.get("sections", []):
        heading = str(section.get("heading", "")).strip()
        if heading:
            document.add_heading(heading, level=min(3, int(section.get("level", 1))))
        for paragraph in section.get("paragraphs", []):
            add_paragraph(document, str(paragraph), bold_lead=True)
        for table in section.get("tables", []):
            add_table(document, table)
        if section.get("page_break"):
            document.add_page_break()
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def register_pdf_font() -> tuple[str, str]:
    candidates = (
        (Path("/Library/Fonts/Times New Roman.ttf"), Path("/Library/Fonts/Times New Roman Bold.ttf")),
        (Path("/System/Library/Fonts/Supplemental/Times New Roman.ttf"), Path("/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf")),
    )
    for regular, bold in candidates:
        if regular.is_file() and bold.is_file():
            if "AktenTimes" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("AktenTimes", str(regular)))
                pdfmetrics.registerFont(TTFont("AktenTimes-Bold", str(bold)))
            return "AktenTimes", "AktenTimes-Bold"
    return "Times-Roman", "Times-Bold"


PDF_FONT, PDF_FONT_BOLD = register_pdf_font()


class InvariantCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        kwargs["invariant"] = 1
        super().__init__(*args, **kwargs)


def pdf_header_footer(case: dict[str, Any]):
    def draw(canv, doc):
        canv.saveState()
        canv.setFont(PDF_FONT, 8)
        canv.setFillColor(colors.HexColor("#5D6871"))
        canv.drawString(2.1 * cm, A4[1] - 1.25 * cm, case["short_title"])
        canv.drawRightString(A4[0] - 2.1 * cm, 1.15 * cm, f"Seite {doc.page}")
        canv.restoreState()
    return draw


def escape_pdf(text: Any) -> str:
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(path: Path, case: dict[str, Any], spec: dict[str, Any]) -> None:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "AktenTitle", parent=styles["Title"], fontName=PDF_FONT_BOLD,
        fontSize=18, leading=22, textColor=colors.HexColor("#183A55"),
        alignment=TA_LEFT, spaceAfter=14,
    )
    heading_style = ParagraphStyle(
        "AktenHeading", parent=styles["Heading2"], fontName=PDF_FONT_BOLD,
        fontSize=13, leading=16, textColor=colors.HexColor("#183A55"),
        spaceBefore=10, spaceAfter=5,
    )
    body_style = ParagraphStyle(
        "AktenBody", parent=styles["BodyText"], fontName=PDF_FONT,
        fontSize=10.5, leading=14.2, textColor=colors.black,
        alignment=TA_LEFT, spaceAfter=7,
    )
    table_style = ParagraphStyle("AktenTable", parent=body_style, fontSize=8.3, leading=10.5)
    story: list[Any] = [Paragraph(escape_pdf(spec["title"]), title_style)]
    for paragraph in spec.get("paragraphs", []):
        story.append(Paragraph(escape_pdf(paragraph), body_style))
    for table in spec.get("tables", []):
        story.extend(pdf_table(table, table_style))
    for section in spec.get("sections", []):
        heading = str(section.get("heading", "")).strip()
        if heading:
            story.append(Paragraph(escape_pdf(heading), heading_style))
        for paragraph in section.get("paragraphs", []):
            story.append(Paragraph(escape_pdf(paragraph), body_style))
        for table in section.get("tables", []):
            story.extend(pdf_table(table, table_style))
        if section.get("page_break"):
            story.append(PageBreak())
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=2.1 * cm, rightMargin=2.1 * cm,
        topMargin=1.8 * cm, bottomMargin=1.65 * cm,
        title=spec["title"], author="Fiktive Aktenstelle", subject=case["short_title"],
    )
    doc.build(story, onFirstPage=pdf_header_footer(case), onLaterPages=pdf_header_footer(case), canvasmaker=InvariantCanvas)


def pdf_table(table_spec: dict[str, Any], style: ParagraphStyle) -> list[Any]:
    headers = table_spec.get("headers", [])
    rows = table_spec.get("rows", [])
    if not headers:
        return []
    matrix = [[Paragraph(f"<b>{escape_pdf(value)}</b>", style) for value in headers]]
    for raw_row in rows:
        values = list(raw_row)
        matrix.append([
            Paragraph(escape_pdf(values[index] if index < len(values) else ""), style)
            for index in range(len(headers))
        ])
    width = A4[0] - 4.2 * cm
    weights = table_spec.get("widths") or [1] * len(headers)
    total = sum(float(value) for value in weights) or len(headers)
    widths = [width * float(value) / total for value in weights]
    table = Table(matrix, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6ED")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#AEB8C0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return [table, Spacer(1, 8)]


def build_csv(path: Path, spec: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter=";")
        headers = spec.get("headers", [])
        if headers:
            writer.writerow(headers)
        writer.writerows(spec.get("rows", []))


def build_eml(path: Path, folder: Path, spec: dict[str, Any]) -> None:
    message = EmailMessage(policy=SMTP)
    message["From"] = spec["from"]
    message["To"] = ", ".join(spec.get("to", []))
    if spec.get("cc"):
        message["Cc"] = ", ".join(spec["cc"])
    raw_date = spec["date"]
    if isinstance(raw_date, datetime):
        message_date = raw_date
    else:
        try:
            message_date = datetime.fromisoformat(str(raw_date))
        except ValueError:
            message_date = parsedate_to_datetime(str(raw_date))
    if message_date.tzinfo is None:
        raise ValueError(f"{path.name}: Date braucht eine Zeitzone: {raw_date!r}")
    message["Date"] = format_datetime(message_date)
    message["Subject"] = spec["subject"]
    message_id = spec.get("message_id")
    if not message_id:
        relative = path.relative_to(folder).as_posix()
        digest = hashlib.sha256(f"{folder.name}/{relative}".encode("utf-8")).hexdigest()[:24]
        message_id = f"<ipakte.{digest}@akten.example>"
    message["Message-ID"] = message_id
    raw_body = spec.get("body", [])
    if isinstance(raw_body, str):
        body = raw_body
    else:
        body = "\n\n".join(str(value) for value in raw_body)
    message.set_content(body, subtype="plain", charset="utf-8")
    for attachment in spec.get("attachments", []):
        attachment_path = folder / safe_rel(attachment)
        if not attachment_path.is_file():
            raise FileNotFoundError(f"{path.name}: Anlage fehlt: {attachment}")
        if attachment_path.suffix.lower() == ".eml":
            nested = BytesParser(policy=SMTP).parsebytes(attachment_path.read_bytes())
            message.add_attachment(nested, filename=attachment_path.name)
            continue
        mime, _ = mimetypes.guess_type(attachment_path.name)
        maintype, subtype = (mime or "application/octet-stream").split("/", 1)
        message.add_attachment(attachment_path.read_bytes(), maintype=maintype, subtype=subtype, filename=attachment_path.name)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(message.as_bytes(policy=SMTP))


def build_eml_batch(eml_specs: list[tuple[Path, Path, dict[str, Any]]]) -> None:
    """Baut E-Mails in Abhängigkeitsreihenfolge und bettet EML-Anlagen echt ein."""
    specs_by_path = {path.resolve(): (path, folder, document) for path, folder, document in eml_specs}
    built: set[Path] = set()
    active: set[Path] = set()

    def build_mail(path_key: Path) -> None:
        if path_key in built:
            return
        if path_key in active:
            raise ValueError(f"Zyklische EML-Anlagenbeziehung: {path_key}")
        active.add(path_key)
        path, folder, document = specs_by_path[path_key]
        for attachment in document.get("attachments", []):
            if Path(attachment).suffix.lower() != ".eml":
                continue
            dependency = (folder / safe_rel(attachment)).resolve()
            if dependency not in specs_by_path:
                raise FileNotFoundError(f"{path.name}: EML-Anlage ohne Quelldatensatz: {attachment}")
            build_mail(dependency)
        build_eml(path, folder, document)
        active.remove(path_key)
        built.add(path_key)

    for path_key in specs_by_path:
        build_mail(path_key)


def write_readme(folder: Path, case: dict[str, Any]) -> None:
    slug = case["slug"]
    parties = "\n".join(f"| {p['name']} | {p['role']} | {p['address']} |" for p in case["parties"])
    questions = "\n".join(f"{index}. {question}" for index, question in enumerate(case["questions"], start=1))
    ambiguities = "\n".join(f"- {item}" for item in case["ambiguities"])
    formats = {}
    for document in case["documents"]:
        formats[document["kind"]] = formats.get(document["kind"], 0) + 1
    format_text = ", ".join(f"{count} {kind.upper()}" for kind, count in sorted(formats.items()))
    text = f"""# {case['title']}

**Plugin:** `{case['plugin']}`
**Aktenordner:** `testakten/{slug}/`

<!-- reserved-example-contacts -->

<!-- BEGIN gesamt-pdf-section (autogen) -->
## Akte komplett herunterladen

[Testakten-Übersicht](../README.md) · [Repository-Start](../../README.md) · [Plugin](../../fachanwalt-urheber-medienrecht/README.md)

Die Akte steht in drei getrennten Fassungen bereit. Das Gesamt-PDF ist die Lesefassung. Das Originalformat-ZIP enthält ausschließlich die nativen Aktenstücke und die vorgeschriebene README.txt; das Gesamt-PDF, redaktionelle Dateien und eine Lösung sind darin nicht enthalten. Das Einzel-PDF-ZIP enthält jede Aktenunterlage als eigene PDF-Datei.

{NOTICE_MARKDOWN}

| Was | Format | Quelle |
| --- | --- | --- |
| Gesamt-PDF | PDF | [`gesamt-pdf/{slug}_gesamt.pdf`](gesamt-pdf/{slug}_gesamt.pdf) |
| Akten-ZIP (Originalformate, ohne Gesamt-PDF) | ZIP | [`testakte-{slug}.zip`]({RELEASE_BASE}/testakte-{slug}.zip) |
| Einzel-PDF-ZIP | ZIP | [`testakte-{slug}-einzelpdfs.zip`]({RELEASE_BASE}/testakte-{slug}-einzelpdfs.zip) |

English: The original-format ZIP contains native working files only. It does not include the combined PDF, editorial files, or a solution. The individual-PDF ZIP contains one PDF per source document.

<!-- END gesamt-pdf-section (autogen) -->

## 1. Auftrag

{case['scenario']}

Erstellen Sie aus den Aktenstücken ein ausführliches Gutachten zur Rechtslage und zu den rechtlich möglichen Reaktionen des jeweiligen IP-Sicherheitennehmers. Trennen Sie belegte Tatsachen, offene Tatsachen, Rechtsfragen, Eilmaßnahmen, Hauptsacheoptionen, Beweisfragen und wirtschaftliche Handlungsalternativen. Die Akte selbst enthält keine Musterlösung.

Rekonstruieren Sie dabei die schuldrechtliche und tatsächliche Kette vollständig: Darlehensvertrag, nicht akzessorische Sicherungsabrede und Sicherungszweck, Auszahlungsvoraussetzungen, Abrufe, Valutierung auf dem Schuldnerkonto, Tilgungen, Zinsen, Kosten und offener Saldo. Eine bloße Bezeichnung als „Sicherheit“ ersetzt weder den Nachweis der gesicherten Hauptschuld noch den Vollzug der Auszahlung.

## 2. Beteiligte

| Beteiligte | Rolle | Anschrift |
| --- | --- | --- |
{parties}

## 3. Konkrete Prüfaufträge

{questions}

## 4. Offene und widersprüchliche Punkte

{ambiguities}

## 5. Aktenumfang

Die Arbeitsakte enthält {len(case['documents'])} Originaldateien: {format_text}. E-Mail-Adressen verwenden ausschließlich reservierte `.example`-Domains. Namen, Unternehmen, Registerdaten und Vorgänge sind erfunden.

## 6. Didaktischer Endpunkt

Der Insolvenzverwalter hat den jeweiligen Erwerbsvertrag bereits unterzeichnet und die Verwertung beziehungsweise den Vollzug angekündigt. Der Erwerber rechnet mit der Übernahme. Eine anwaltliche Reaktion, ein gerichtlicher Antrag, eine Genehmigung oder eine abschließende Vereinbarung liegt noch nicht vor.
"""
    (folder / "README.md").write_text(text, encoding="utf-8")


def write_rubric(folder: Path, case: dict[str, Any]) -> None:
    data = {
        "schema_version": 1,
        "id": case["slug"],
        "title": case["title"],
        "plugin": case["plugin"],
        "task": "Gutachten zur Rechtslage und zu möglichen Reaktionen des IP-Sicherheitennehmers",
        "artifacts": {
            "combined_pdf": f"gesamt-pdf/{case['slug']}_gesamt.pdf",
            "original_zip": f"testakte-{case['slug']}.zip",
            "individual_pdf_zip": f"testakte-{case['slug']}-einzelpdfs.zip",
        },
        "criteria": [
            "Sachen, Forderungen und sonstige Rechte werden gegenstandsbezogen getrennt.",
            "Rechtekette und Sicherheitenbestellung werden aus Primärakten rekonstruiert.",
            "Hauptschuld, Sicherungszweck, Auszahlungsvoraussetzungen, Abrufe, tatsächliche Valutierung, Tilgungen und offener Saldo werden aus getrennten Primärbelegen zu einer prüfbaren Kette verbunden.",
            "Der unterschriebene Erwerbsvertrag wird nach Vollzugsstand und betroffenen Rechten ausgewertet.",
            "Eilbedarf, Parteistellung, Anträge, Beweisangebote und wirtschaftliche Alternativen werden geprüft.",
            "Kein Aktenstück wird als bereits erstellte anwaltliche Lösung missverstanden.",
        ],
        "must_not_contain": ["Musterlösung", "fertiges Gutachten", "abschließende Anspruchsstrategie"],
    }
    (folder / "rubric.yaml").write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def write_xlsx_specs(spec_path: Path, workbooks: list[dict[str, Any]]) -> None:
    spec_path.parent.mkdir(parents=True, exist_ok=True)
    spec_path.write_text(json.dumps({"workbooks": workbooks}, ensure_ascii=False, indent=2), encoding="utf-8")


def run_workbook_builder(spec_path: Path, preview_dir: Path) -> None:
    default_runtime = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies"
    runtime = Path(os.environ.get("CODEX_PRIMARY_RUNTIME", default_runtime))
    node = runtime / "node" / "bin" / "node"
    packages = runtime / "node" / "node_modules"
    if not node.is_file() or not packages.is_dir():
        raise RuntimeError("Gebündelte Node-Arbeitsumgebung fehlt")
    with tempfile.TemporaryDirectory(prefix="urheberrecht-xlsx-") as raw:
        temp = Path(raw)
        (temp / "node_modules").symlink_to(packages, target_is_directory=True)
        runner = temp / "build.mjs"
        shutil.copy2(ROOT / "scripts" / "build-urheberrecht-ip-insolvenz-workbooks.mjs", runner)
        subprocess.run([str(node), str(runner), str(spec_path), str(preview_dir)], cwd=temp, check=True)


def build_source_files(selected: set[str] | None = None) -> list[dict[str, Any]]:
    selected_cases = [case for case in cases() if selected is None or case["slug"] in selected]
    if selected and len(selected_cases) != len(selected):
        known = {case["slug"] for case in cases()}
        raise SystemExit(f"Unbekannte Slugs: {sorted(selected - known)}")
    workbook_specs: list[dict[str, Any]] = []
    eml_specs: list[tuple[Path, Path, dict[str, Any]]] = []
    for case in selected_cases:
        folder = clear_case(case)
        write_readme(folder, case)
        write_rubric(folder, case)
        for document in case["documents"]:
            rel = safe_rel(document["path"])
            path = folder / rel
            kind = document["kind"]
            if path.suffix.lower() != f".{kind}":
                raise ValueError(f"{case['slug']}: Dateiendung passt nicht zu {kind}: {rel}")
            if kind == "docx":
                build_docx(path, case, document)
            elif kind == "pdf":
                build_pdf(path, case, document)
            elif kind == "csv":
                build_csv(path, document)
            elif kind == "xlsx":
                workbook_specs.append({"path": str(path), "title": document["title"], "sheets": document["sheets"]})
            elif kind == "eml":
                eml_specs.append((path, folder, document))
            else:
                raise ValueError(f"{case['slug']}: nicht unterstützter Dokumenttyp {kind!r}")
    build_root = ROOT / "tmp" / "urheberrecht-ip-insolvenz"
    spec_path = build_root / "workbooks.json"
    preview_dir = build_root / "xlsx-previews"
    write_xlsx_specs(spec_path, workbook_specs)
    if preview_dir.exists():
        shutil.rmtree(preview_dir)
    preview_dir.mkdir(parents=True)
    run_workbook_builder(spec_path, preview_dir)
    # EML-Anlagen müssen als echte message/rfc822-Teile eingebettet sein. Die
    # Abhängigkeiten werden deshalb vor der jeweils übergeordneten Nachricht
    # gebaut; Zyklen sind ein Quelldatenfehler.
    build_eml_batch(eml_specs)
    return selected_cases


def build_combined_pdfs(selected_cases: list[dict[str, Any]]) -> None:
    command = [sys.executable, str(ROOT / "scripts" / "build-testakte-gesamt-pdf.py")]
    command.extend(case["slug"] for case in selected_cases)
    subprocess.run(command, cwd=ROOT, check=True)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def spec_sha256(document: dict[str, Any]) -> str:
    payload = json.dumps(
        document,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def write_artifact_manifest(selected_cases: list[dict[str, Any]]) -> None:
    """Bindet jede Quelldefinition an das eingecheckte Binärartefakt."""
    if ARTIFACT_MANIFEST.is_file():
        manifest = json.loads(ARTIFACT_MANIFEST.read_text(encoding="utf-8"))
    else:
        manifest = {"schema_version": 1, "cases": {}}
    if manifest.get("schema_version") != 1 or not isinstance(manifest.get("cases"), dict):
        raise ValueError(f"Ungültiges Artefaktmanifest: {ARTIFACT_MANIFEST}")
    for case in selected_cases:
        folder = TESTAKTEN / case["slug"]
        documents = []
        for document in case["documents"]:
            artifact = folder / safe_rel(document["path"])
            if not artifact.is_file():
                raise FileNotFoundError(artifact)
            documents.append({
                "path": document["path"],
                "spec_sha256": spec_sha256(document),
                "artifact_sha256": sha256_file(artifact),
            })
        combined = folder / "gesamt-pdf" / f"{case['slug']}_gesamt.pdf"
        if not combined.is_file():
            raise FileNotFoundError(combined)
        manifest["cases"][case["slug"]] = {
            "documents": documents,
            "combined_pdf_sha256": sha256_file(combined),
        }
    ARTIFACT_MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    ARTIFACT_MANIFEST.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("slugs", nargs="*")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    selected = set(args.slugs) or None
    selected_cases = build_source_files(selected)
    build_combined_pdfs(selected_cases)
    write_artifact_manifest(selected_cases)
    counts = {case["slug"]: len(case["documents"]) for case in selected_cases}
    print(json.dumps({"built": counts}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
