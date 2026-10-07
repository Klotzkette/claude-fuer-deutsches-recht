#!/usr/bin/env python3
"""Prüft Arbeitsmappen, absichtliche Abweichungen und vollständigen PDF-Export."""

import re
from datetime import datetime, time
import unittest
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader
from bau_rundum_excel import ROOT, RELEASE, specifications

SPECS = specifications()
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def formula_text(value):
    # Office may omit optional quotes around simple worksheet names.
    return re.sub(r"'([^\W\d]\w*)'!", r"\1!", value)


class ExcelUnterlagen(unittest.TestCase):
    def test_ten_distinct_case_workbooks(self):
        self.assertEqual(len(SPECS), 10)
        self.assertEqual(len({s["slug"] for s in SPECS}), 10)
        for spec in SPECS:
            self.assertRegex(spec["filename"], r"^[0-9]{2}_[A-Za-z0-9_]+\.xlsx$")
            self.assertGreaterEqual(sum(len(s["rows"]) for s in spec["sheets"]), 60)
            self.assertEqual(len(spec["intentional"]), 2)
            self.assertGreaterEqual(len(spec["controls"]), 3)

    def test_saved_sources_formulas_and_cached_values(self):
        for spec in SPECS:
            path = ROOT / "testakten" / spec["slug"] / spec["filename"]
            wf = load_workbook(path, data_only=False)
            wv = load_workbook(path, data_only=True)
            self.assertEqual(wf.sheetnames, [s["name"] for s in spec["sheets"]])
            for s in spec["sheets"]:
                with self.subTest(case=spec["slug"], sheet=s["name"]):
                    sheet = wf[s["name"]]
                    self.assertEqual(sheet.sheet_state, "visible")
                    self.assertFalse(any(d.hidden for d in sheet.row_dimensions.values()))
                    self.assertFalse(any(d.hidden for d in sheet.column_dimensions.values()))
                    self.assertEqual(sheet.freeze_panes, "A9")
                    self.assertEqual(len(sheet.tables), 1)
                    self.assertFalse(sheet.sheet_view.showGridLines)
                    self.assertEqual(sheet.print_title_rows.replace('$', ''), '8:8')
                    self.assertEqual(sheet.page_setup.fitToHeight, 0)
                    self.assertEqual(sheet.page_setup.fitToWidth, 1)
                    self.assertEqual(sheet.max_row, 8 + len(s["rows"]))
                    for cells in sheet:
                        for cell in cells:
                            if cell.data_type == 'f':
                                self.assertIsNotNone(wv[s['name']][cell.coordinate].value)
                                self.assertNotEqual(wv[s['name']][cell.coordinate].data_type, 'e')
                    for i, row in enumerate(s["rows"], 9):
                        for j, expected in enumerate(row, 1):
                            cell = sheet.cell(i, j)
                            actual = cell.value
                            if expected is None or expected == '':
                                self.assertIn(actual, (None, ''))
                            elif isinstance(expected, str) and expected.startswith('='):
                                self.assertEqual(formula_text(actual), formula_text(expected))
                                cached = wv[s["name"]].cell(i, j)
                                self.assertIsNotNone(cached.value, (path, cell.coordinate))
                                self.assertNotEqual(cached.data_type, 'e', (path, cell.coordinate, cached.value))
                            elif s["columns"][j-1]["type"] == 'date':
                                self.assertEqual(actual.strftime('%Y-%m-%d'), expected)
                            elif s["columns"][j-1]["type"] == 'time':
                                self.assertEqual(actual.strftime('%H:%M'), expected)
                            else:
                                self.assertEqual(actual, expected, (path, cell.coordinate))
                                if isinstance(actual, (int, float)) and actual != 0 and abs(actual) < 0.005 and s['columns'][j-1]['type'] == 'number':
                                    self.assertIn('E+', cell.number_format, 'Kleine Messwerte dürfen nicht als Null erscheinen.')
                                if s["columns"][j-1]["type"] == 'text':
                                    self.assertIsInstance(actual, str)
                    for control in spec["controls"]:
                        actual = wv[control["sheet"]][control["cell"]].value
                        if isinstance(control["value"], (int, float)):
                            self.assertAlmostEqual(actual, control["value"], places=5)
                        else:
                            self.assertEqual(actual, control["value"])
            wf.close()
            wv.close()

    def test_intentional_discrepancies_remain_specific(self):
        for spec in SPECS:
            book = load_workbook(ROOT / "testakten" / spec["slug"] / spec["filename"], data_only=False)
            for item in spec["intentional"]:
                with self.subTest(case=spec["slug"], cell=item["cell"]):
                    actual = book[item["sheet"]][item["cell"]].value
                    if isinstance(actual, datetime):
                        actual = actual.strftime('%Y-%m-%d')
                    elif isinstance(actual, time):
                        actual = actual.strftime('%H:%M')
                    self.assertEqual(actual, item["actual"])
                    self.assertNotEqual(item["actual"], item["expected"])
                    self.assertGreater(len(item["reason"]), 25)
                    self.assertTrue(item["source"])
            book.close()

    def test_no_external_links_macros_hidden_solutions_or_error_cells(self):
        for spec in SPECS:
            path = ROOT / "testakten" / spec["slug"] / spec["filename"]
            with zipfile.ZipFile(path) as archive:
                self.assertFalse(any('externallinks' in name.lower() or 'vbaproject' in name.lower() for name in archive.namelist()))
                for name in archive.namelist():
                    if not name.endswith('.xml'):
                        continue
                    text = archive.read(name).decode('utf-8')
                    self.assertNotRegex(text, r'Musterlösung|Lösungsmatrix|absichtlich eingebaut|Fehlerversteck|Red.Team|Trainingsfehler')
                    self.assertNotIn('§', text)
                    if name.startswith('xl/worksheets/sheet'):
                        document = ET.fromstring(text)
                        self.assertEqual(document.findall('.//s:c[@t="e"]', NS), [])
                        self.assertEqual(document.findall('.//s:f[@t="array"]', NS), [])

    def test_new_workbook_rows_survive_pdf_export_and_are_linked(self):
        for spec in SPECS:
            folder = ROOT / "testakten" / spec["slug"]
            readme = (folder / "README.md").read_text(encoding='utf-8')
            self.assertIn(f"]({spec['filename']})", readme)
            inventory = re.search(r"^\| (?:Datei \| Format|Nr\. \| Datei \| Datum \| Inhalt) \|\n(?:\|[^\n]*\n)+", readme, re.MULTILINE)
            self.assertIsNotNone(inventory)
            self.assertIn(f"]({spec['filename']})", inventory[0])
            headings = re.findall(r"^## (1\.\d+)\. ", readme, re.MULTILINE)
            self.assertEqual(len(headings), len(set(headings)))
            self.assertEqual(set(re.findall(r"bauwirtschaft-rundum-v\d+\.\d+\.\d+", readme)), {RELEASE})
            pdf = PdfReader(folder / "gesamt-pdf" / f"{spec['slug']}_gesamt.pdf")
            text = re.sub(r'\s+', '', '\n'.join(p.extract_text() or '' for p in pdf.pages))
            for sheet in spec["sheets"]:
                self.assertIn(re.sub(r'\s+', '', sheet['title']), text)
                for index, record in enumerate(sheet['rows']):
                    for j, value in enumerate(record):
                        if sheet['columns'][j]['type'] == 'text' and isinstance(value, str) and not value.startswith('=') and len(value) >= 3:
                            self.assertIn(re.sub(r'\s+', '', value), text, (spec['slug'], sheet['name'], index, value))


if __name__ == '__main__':
    unittest.main()
