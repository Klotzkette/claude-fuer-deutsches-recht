#!/usr/bin/env python3
"""Einbecker Originalakte. Autor: Klotzkette. Keine zentralen Schreibzugriffe."""
from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import mimetypes
import os
from datetime import datetime
from email import policy
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, PngImagePlugin
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from akten_build_runtime import serif_font_path, screen_font

ROOT = Path(__file__).resolve().parents[1]
SLUG = "bauwirtschaft-hoai-buergerhaus-einbeck"
CASE = ROOT / "testakten" / SLUG
ASSETS = Path(os.environ.get("EINBECK_ASSETS", "/tmp/bauwirtschaft-assets"))
AUTHOR = "Klotzkette"
FILES = {
 1: "01_2020-11-12_Bedarfsbeschluss.docx", 2: "02_2021-01-18_Architektenvertrag.pdf",
 3: "03_2021-02-04_Grundlagengespraech.docx", 4: "04_2021-02-10_Bestandsaufnahme.pdf",
 5: "05_2021-03-18_Variantenuntersuchung.pdf", 6: "06_2021-03-25_Variantenentscheidung.eml",
 7: "07_2021-06-10_Entwurfsplan.pdf", 8: "08_2021-06-11_Objektbeschreibung.docx",
 9: "09_2021-06-14_Kostenberechnung.xlsx", 10: "10_2021-06-15_Fachplanerabstimmung.pdf",
 11: "11_2022-01-20_Bauantrag.pdf", 12: "12_2022-01-20_Lageplan.pdf",
 13: "13_2022-01-20_Bauzeichnungen.pdf", 14: "14_2022-01-20_Baubeschreibung.docx",
 15: "15_2022-02-03_Nachforderung.eml", 16: "16_2022-02-14_Brandschutzergänzung.pdf",
 17: "17_2022-03-17_Baugenehmigung.pdf", 18: "18_2022-04-28_Freigabe.eml",
 19: "19_2022-04-25_Rahmenterminplan.xlsx", 20: "20_2022-06-02_Ausführungsplan.pdf",
 21: "21_2022-06-06_Rampenanschluss.pdf", 22: "22_2023-03-07_Planstandportal.png",
 23: "23_2023-03-08_Leistungsverzeichnis.docx", 24: "24_2023-03-08_Bepreistes_LV.xlsx",
 25: "25_2023-04-04_Angebot_Leinebau.pdf", 26: "26_2023-04-04_Angebot_Hagedorn.pdf",
 27: "27_2023-04-11_Angebotsprüfung.pdf", 28: "28_2023-04-18_Bauvertrag_Zuschlag.docx",
 29: "29_2023-08-14_Vereinsraumnutzung.eml", 30: "30_2024-03-18_Bautagebuch.csv",
 31: "31_2023-09-05_Behinderungsanzeige.pdf", 32: "32_2023-09-06_Bauleiterchat.png",
 33: "33_2023-09-12_Nachtragsangebot.pdf", 34: "34_2023-09-15_Nachtragsauftrag.eml",
 35: "35_2024-02-16_Gemeinsames_Aufmaß.pdf", 36: "36_2024-03-01_Schlussrechnung.pdf",
 37: "37_2024-04-08_Kostenfeststellung.xlsx", 38: "38_2024-03-18_Abnahme.pdf",
 39: "39_2024-04-05_Übergabe.docx", 40: "40_2024-04-08_Archivverzeichnis.txt",
 41: "41_2026-09-08_Mängelmeldung.eml", 42: "42_2026-09-07_Feuchte_Saalnordwand.png",
 43: "43_2026-09-17_Objektbegehung.pdf", 44: "44_2026-09-21_Gewährleistungsfeststellung.docx",
}
PHASES = {1:[1,2,3,4],2:[5,6],3:[7,8,9,10],4:[11,12,13,14,15,16,17],5:[18,19,20,21,22],6:[23,24],7:[25,26,27,28],8:[29,30,31,32,33,34,35,36,37,38,39,40],9:[41,42,43,44]}
ACTORS = {
 "verein": ("Bürgerhaus Leinewinkel e.V.", "Lindengasse 18 · 37574 Einbeck", "vorstand@buergerhaus-leinewinkel.de", "Maren Döring, Vorsitzende"),
 "plan": ("Hartwig Architektur", "Mühlenweg 7 · 37574 Einbeck", "projekt@hartwig-architektur.de", "Lena Hartwig, Architektin"),
 "bau": ("Leinebau Hochbau GmbH", "Am Werkhof 6 · 37574 Einbeck", "bauleitung@leinebau-hochbau.de", "Torsten Albrecht, Bauleiter"),
 "hage": ("Hagedorn Bauhandwerk GmbH", "Am Gewerbering 12 · 37586 Dassel", "angebot@hagedorn-bauhandwerk.de", "Jens Hagedorn, Geschäftsführer"),
 "amt": ("Stadt Einbeck · Bauaufsicht", "Teichenweg 1 · 37574 Einbeck", "bauaufsicht@einbeck.de", "Klara Brandes, im Auftrag"),
 "technik": ("Ingenieurbüro Mertens", "Brückenstraße 9 · 37154 Northeim", "planung@mertens-haustechnik.de", "Paul Mertens, Fachplaner"),
}
QTY = [1,216,60,24,80,95,36,144,80,3,20.25,9]
PLAN_EP = [8500,38,85,310,145,128,185,48,245,2450,240,155]
LEINE_EP = [9000,41,92,330,152,135,205,52,268,2590,260,172]
HAGE_EP = [8200,45,104,315,168,142,175,58,285,2710,245,190]
FINAL_QTY = [1,216,66,25.2,80,98,42,150,80,3,21.45,9.4]
UNITS = ["psch","m²","m³","m³","m²","m²","m²","m²","m²","St","m²","m"]
ITEMS = [
 ("Baustelleneinrichtung", "Einrichten und Räumen für Los 1 einschließlich Bauzaun, Baustromverteiler und Sanitärcontainer. Vorhaltung vom 21.08.2023 bis 15.03.2024. Betriebskosten eigener Geräte sind enthalten."),
 ("Selektiver Rückbau", "Aufnehmen der alten Bodenbeläge und nichttragenden Einbauten im 216 m² großen Bestand. Staubschutz zum südlichen Geräteanbau und sortenreine Entsorgung nicht gefährlicher Stoffe sind enthalten. Verdeckte Schadstoffe werden vor Bearbeitung angezeigt."),
 ("Fundamentaushub", "Aushub der Streifenfundamente und Anschlussgräben bis 1,20 m unter Gelände einschließlich seitlicher Lagerung. Leitungen dürfen nicht durchtrennt werden. Handschachtung an vorher georteten Querungen ist im Einheitspreis enthalten; Umlegungen sind nicht enthalten."),
 ("Streifenfundamente", "Betonfundamente nach Tragwerksplan T-02 mit 50 cm Breite und 80 cm Höhe einschließlich Schalung und Bewehrung gemäß Stahlliste, insgesamt 24 m³. Abrechnung nach tatsächlich eingebautem geometrischem Volumen."),
 ("Bodenplatte Anbau", "Gedämmte Bodenplatte für 80 m² Anbau einschließlich 20 cm Stahlbeton, 12 cm druckfester Dämmung und Sauberkeitsschicht. Anschlussfuge an den Bestand nicht starr überbrücken. Oberkante Rohplatte bei minus 0,16 m."),
 ("Außenmauerwerk Anbau", "24 cm tragendes Mauerwerk mit außenliegender 16 cm Mineralwolldämmung und mineralischem Putzsystem. Fensteröffnungen werden vollständig abgezogen. Stürze und Anschlüsse sind enthalten; Metallfenster werden separat vergeben."),
 ("Abdichtung Bestandsockel", "Freilegen, Reinigen und zweilagiges Abdichten der nördlichen Bestandswand im Anschlussbereich. Bearbeitung vom Fundamentansatz bis 30 cm über Gelände. Schutzlage und Wiederverfüllung sind enthalten. Die Abdichtung ersetzt keine nachträgliche Horizontalsperre."),
 ("Innenputz Bestand", "Lose Putzbereiche entfernen und mineralischen Neuputz herstellen, mittlere Dicke 20 mm, einschließlich Untergrundvorbereitung. Beschichtungsaufbau diffusionsoffen; Untergrundfeuchte vor Beginn gemeinsam dokumentieren. Eine chemische Horizontalsperre ist nicht enthalten."),
 ("Dach Anbau", "Flachdachaufbau einschließlich Holztragwerk nach T-02, Dampfbremse, Gefälledämmung im Mittel 20 cm, zweilagiger Abdichtung, Attika und Dachentwässerung bis zum Standrohr. Gefälle vom Bestandsanschluss wegführen."),
 ("Außentüren", "Drei thermisch getrennte Aluminiumtüren nach Türliste, lichte Durchgangsbreite jeweils 1,00 m, einschließlich Beschlägen, Dichtungen und Anschlussarbeiten. Hauptzugang mit niveaugleichem Durchgang, Rinne gemäß eigener Position."),
 ("Zugangsrampe", "Stahlbetonrampe, zwei Läufe mit zusammen 12,00 m nutzbarer Länge und 1,50 m nutzbarer Breite. Höhendifferenz 0,60 m, Längsgefälle 5 Prozent. Einschließlich Zwischenpodest 1,50 m mal 1,50 m ergeben sich 20,25 m². Handläufe und Radabweiser sind enthalten."),
 ("Entwässerungsrinne", "Linienrinne mit herausnehmbarem Rost am Eingangs- und Rampenanschluss, lichte Breite 150 mm, Anschluss an den Regenwasserstrang. Gefälle, Bettung und Dichtheitsprobe sind enthalten. Anschlussdetail D-12 ist Vertragsgrundlage."),
]
DATA = {}

def euro(v):
    return f"{v:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")

def register_fonts():
    for name, bold in [("TNR",False),("TNRB",True)]:
        pdfmetrics.registerFont(TTFont(name, str(serif_font_path(bold=bold))))
    pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNRB", italic="TNR", boldItalic="TNRB")

BODY = ParagraphStyle("Text",fontName="TNR",fontSize=11,leading=14,spaceAfter=7)
SMALL = ParagraphStyle("Klein",parent=BODY,fontSize=9,leading=11,spaceAfter=5)
HEAD = ParagraphStyle("Abschnitt",parent=BODY,fontName="TNRB",spaceBefore=9,spaceAfter=10,keepWithNext=True)
TITLE = ParagraphStyle("Titel",parent=HEAD,fontSize=16,leading=19,spaceAfter=14)

def para(s,style=BODY):
    return Paragraph(escape(str(s)).replace("\n","<br/>"),style)

def record(n, issuer, recipient, title, sections, signer=None, attachments=()):
    DATA[n] = dict(number=n,file=FILES[n],date=FILES[n][3:13],issuer=issuer,recipient=recipient,title=title,
                   sections=sections,signer=signer or ACTORS[issuer][3],attachments=list(attachments))

def table(rows, widths=None):
    if widths is None:
        widths=[495/len(rows[0])]*len(rows[0])
    t=Table([[para(c,SMALL) for c in r] for r in rows],colWidths=widths,repeatRows=1,hAlign="LEFT")
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e8ecee")),("VALIGN",(0,0),(-1,-1),"TOP"),
                          ("LINEBELOW",(0,0),(-1,0),.6,colors.grey),("LINEBELOW",(0,1),(-1,-1),.25,colors.lightgrey),
                          ("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),
                          ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
    return t

def letter_pdf(d):
    issuer=ACTORS[d["issuer"]]
    def frame(c,doc):
        c.setFont("TNRB",14); c.drawString(50,800,issuer[0])
        c.setFont("TNR",9); c.drawString(50,785,issuer[1]); c.drawString(50,772,issuer[2])
        c.setFont("TNR",8); c.drawString(50,32,f"BH-21-04 · {d['title']}"); c.drawRightString(545,32,f"Seite {doc.page}")
    story=[para(ACTORS[d["recipient"]][0]+"\n"+ACTORS[d["recipient"]][1],SMALL),
           para(f"Einbeck, {datetime.fromisoformat(d['date']).strftime('%d.%m.%Y')} · Bezug BH-21-04",SMALL),
           para(d["title"],TITLE),para("Sehr geehrte Damen und Herren,")]
    for s in d["sections"]:
        if isinstance(s,str): story.append(para(s))
        elif s[0]=="h": story.append(para(s[1],HEAD))
        elif s[0]=="t": story += [table(s[1],s[2] if len(s)>2 else None),Spacer(1,9)]
        elif s[0]=="page": story.append(PageBreak())
    story.append(KeepTogether([Spacer(1,7),para("Mit freundlichen Grüßen"),para("gez. "+d["signer"]),
        para("Anlagen: "+("; ".join(FILES[x] for x in d["attachments"]) or "keine"),SMALL)]))
    SimpleDocTemplate(str(CASE/d["file"]),pagesize=A4,leftMargin=50,rightMargin=50,topMargin=90,bottomMargin=53,
                      title=d["title"],author=AUTHOR).build(story,onFirstPage=frame,onLaterPages=frame)

def letter_docx(d):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(2.2); sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    sec.footer_distance=Cm(.8)
    for s in doc.styles:
        if s.type==1:
            s.font.name="Times New Roman";s.font.size=Pt(11);s.font.color.rgb=RGBColor(0,0,0)
            s.paragraph_format.space_after=Pt(7)
            fonts=s.element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if key.endswith("Theme"):del fonts.attrib[key]
            for key in ("ascii","hAnsi","eastAsia","cs"):fonts.set(qn("w:"+key),"Times New Roman")
            for border in s.element.xpath("./w:pPr/w:pBdr"):border.getparent().remove(border)
    doc.styles["Title"].font.size=Pt(16)
    doc.styles["Heading 1"].font.size=Pt(12);doc.styles["Heading 1"].font.bold=True
    doc.styles["Heading 1"].paragraph_format.space_before=Pt(9)
    doc.styles["Heading 1"].paragraph_format.space_after=Pt(10)
    props=doc.core_properties;props.author=AUTHOR;props.last_modified_by=AUTHOR;props.title=d["title"];props.language="de-DE"
    props.created=datetime.fromisoformat(d["date"]);props.modified=props.created
    issuer=ACTORS[d["issuer"]]
    doc.add_paragraph(issuer[0]).runs[0].bold=True
    doc.add_paragraph(issuer[1]+"\n"+issuer[2])
    doc.add_paragraph(ACTORS[d["recipient"]][0]+"\n"+ACTORS[d["recipient"]][1])
    doc.add_paragraph(f"Einbeck, {datetime.fromisoformat(d['date']).strftime('%d.%m.%Y')} · Bezug BH-21-04")
    doc.add_paragraph(d["title"],"Title");doc.add_paragraph("Sehr geehrte Damen und Herren,")
    for s in d["sections"]:
        if isinstance(s,str):doc.add_paragraph(s).paragraph_format.keep_together=True
        elif s[0]=="h":doc.add_paragraph(s[1],"Heading 1")
        elif s[0]=="page":doc.add_page_break()
        elif s[0]=="t":
            t=doc.add_table(rows=0,cols=len(s[1][0]));t.style="Table Grid"
            for ri,row in enumerate(s[1]):
                cells=t.add_row().cells
                t.rows[ri]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for cell,val in zip(cells,row):
                    cell.text=str(val)
                    for p in cell.paragraphs:
                        for run in p.runs:run.font.size=Pt(10);run.bold=ri==0
                if ri==0:
                    prop=t.rows[ri]._tr.get_or_add_trPr();prop.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph()
    doc.add_paragraph("Mit freundlichen Grüßen").paragraph_format.keep_with_next=True
    doc.add_paragraph("gez. "+d["signer"]).paragraph_format.keep_with_next=True
    doc.add_paragraph("Anlagen: "+("; ".join(FILES[x] for x in d["attachments"]) or "keine"))
    footer=sec.footer.paragraphs[0];footer.text="BH-21-04 · Seite "
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    doc.save(CASE/d["file"])

def eml(d):
    sender=ACTORS[d["issuer"]]; recipient=ACTORS[d["recipient"]]
    dt=datetime.fromisoformat(d["date"]+"T09:24:00").replace(tzinfo=ZoneInfo("Europe/Berlin"))
    m=EmailMessage(policy=policy.SMTP)
    m["From"]=f"{sender[3]} <{sender[2]}>";m["To"]=f"{recipient[3]} <{recipient[2]}>"
    m["Date"]=format_datetime(dt);m["Subject"]="BH-21-04 / "+d["title"]
    m["Message-ID"]=f"<bh2104.{d['number']}.{d['date']}@{sender[2].split('@')[1]}>"
    m["Return-Path"]=f"<{sender[2]}>"
    m["Received"]=f"from mail.{sender[2].split('@')[1]} by archiv.buergerhaus-leinewinkel.de with ESMTP; {format_datetime(dt)}"
    m["Content-Language"]="de-DE"
    body="Sehr geehrte Damen und Herren,\n\n"+"\n\n".join(s for s in d["sections"] if isinstance(s,str))
    body+="\n\nMit freundlichen Grüßen\n"+sender[3]+"\n"+sender[0]+"\n"+sender[1]+"\n"+sender[2]
    body+="\n\nAnlagen: "+("; ".join(FILES[x] for x in d["attachments"]) or "keine")+"\n"
    m.set_content(body,charset="utf-8")
    for n in d["attachments"]:
        p=CASE/FILES[n];mime=mimetypes.guess_type(p.name)[0] or "application/octet-stream"
        main,sub=mime.split("/",1);m.add_attachment(p.read_bytes(),maintype=main,subtype=sub,filename=p.name)
    (CASE/d["file"]).write_bytes(m.as_bytes())

def define_documents():
    record(1,"verein","plan","Bedarfsbeschluss zum Bürgerhaus",[
      "der Vorstand hat am 12.11.2020 von 18:30 bis 20:15 Uhr im bisherigen Vereinsraum über den Umbau beraten. Anwesend waren Maren Döring, Kassierer Uwe Feldmann, Schriftführerin Silke Bertram und Hauswart Ralf Winter. Alle vier stimmen dem nachstehenden Raum- und Finanzierungsrahmen zu. Die Mitgliederversammlung hat den Vorstand am 28.10.2020 zur Planungsvorbereitung ermächtigt.",
      ("h","1 Nutzung und Bestand"),"Das eingeschossige Haus aus dem Jahr 1968 steht im Eigentum des Vereins. Der Saal wird für Seniorennachmittage, Chorproben und private Feiern genutzt. Die Belegung soll einschließlich Helfern 120 Personen nicht überschreiten. Im 24 m² großen Vereinsraum sind Besprechungen für höchstens 18 Personen vorgesehen, keine gewerbliche Küche und kein regelmäßiger Sportbetrieb. Der vorhandene Saal mit 120 m² muss erhalten bleiben.",
      "Ein stufenfreier Zugang, ein barrierefrei nutzbares WC, getrennte weitere Sanitärräume und ein beheizbarer Eingangsbereich fehlen. Der alte Eingang liegt rund 60 cm über dem Hof. Bei Regen ist die Nordwand innen fleckig. Der Hauswart erinnert sich an eine Leitung zum inzwischen abgebauten Pumpenschacht, kennt ihren Verlauf jedoch nicht.",
      ("h","2 Flächen und Finanzierung"),("t",[["Bedarf","Nutzfläche","Betriebliche Anforderung"],["Saal im Bestand","120 m²","Bestuhlung für 96 Gäste und bis zu 24 Helfer"],["Vereinsraum","24 m²","18 Sitzplätze, abschließbarer Schrank"],["Anbau","etwa 65 bis 70 m²","Foyer, Sanitär, Lager und Technik"],["Außenbereich","etwa 300 m²","Rampe, Lieferzugang und Fahrradplätze"]]),
      "Der Verein setzt zunächst 430.000,00 EUR einschließlich Umsatzsteuer als verfügbaren Gesamtrahmen an. 180.000,00 EUR stammen aus Rücklagen und zugesagten Spenden, 150.000,00 EUR aus einem zugesagten Vereinsdarlehen und 100.000,00 EUR aus weiteren verbindlich zugesagten privaten Zuwendungen. Es werden keine öffentlichen Fördermittel beantragt. Eigenleistungen sind nur für bewegliche Ausstattung vorgesehen und dürfen die Baukalkulation nicht vermindern.",
      ("h","3 Auftrag zur Vorbereitung"),"Frau Hartwig soll einen Vertragsvorschlag für Gebäudeplanung sowie Vorschläge für Tragwerks- und Haustechnikplanung vorlegen. Vor Öffnung von Wänden ist ein gesonderter Untersuchungstermin abzustimmen. Der Chor kann während des Baus in die benachbarte Gaststätte ausweichen; das Archiv und die vereinseigene Küche werden vor Baubeginn geräumt. Die Fertigstellung wird für Frühjahr 2024 angestrebt. Eine Baustellenöffnung im Winter wird nicht vorausgesetzt."
    ],"Maren Döring, Vorsitzende; Silke Bertram, Schriftführerin")
    record(2,"plan","verein","Architektenvertrag Gebäudeplanung",[
      "Bürgerhaus Leinewinkel e.V., vertreten durch Maren Döring und Uwe Feldmann, beauftragt Hartwig Architektur, Inhaberin Lena Hartwig, mit Umbau und nördlichem Anbau des Bürgerhauses Lindengasse 18 in Einbeck. Die Parteien vereinbaren die folgenden Leistungen und Vergütungen individuell.",
      ("h","1 Leistungsgegenstand"),"Geschuldet sind die Grundleistungen des Leistungsbildes Gebäude für die Leistungsphasen 1 bis 9 nach Anlage 10 zur HOAI. Zunächst werden die Leistungsphasen 1 bis 4 abgerufen. Die Leistungsphasen 5 bis 9 werden erst nach schriftlichem Abruf ausgeführt. Die Gebäudeteile werden als ein funktional zusammenhängendes Objekt geplant. Die Abstimmung mit Tragwerks- und Haustechnikplanung gehört zum Auftrag; deren Fachleistungen sind nicht im Honorar enthalten.",
      "Die vorhandenen Räume und der Bedarf folgen dem Vorstandsprotokoll vom 12.11.2020. Frau Hartwig führt die Ortsbesichtigung durch, stellt Varianten mit Kosten und Terminen gegenüber und klärt den Untersuchungsbedarf. Verdeckte Bauteile werden nicht ohne Zustimmung geöffnet. Bestandsaufmaß, vier Feuchtemessstellen und zwei Suchschlitze sind als zusätzliche Leistungen mit 3.000,00 EUR netto enthalten. Laboruntersuchungen oder eine flächendeckende Leitungssuche werden nur nach gesonderter Beauftragung ausgeführt.",
      ("h","2 Honorar"),("t",[["Leistungsphase","Anteil","Nettohonorar EUR"],["1","2 Prozent","960,00"],["2","7 Prozent","3.360,00"],["3","15 Prozent","7.200,00"],["4","3 Prozent","1.440,00"],["5","25 Prozent","12.000,00"],["6","10 Prozent","4.800,00"],["7","4 Prozent","1.920,00"],["8","32 Prozent","15.360,00"],["9","2 Prozent","960,00"]]),
      "Das individuell vereinbarte Pauschalhonorar für alle Grundleistungen beträgt 48.000,00 EUR netto. Die Prozentanteile dienen der Zuordnung dieser Pauschale zu den abgerufenen Leistungen und sind keine Behauptung einer zwingenden Mindestvergütung. Mit den 3.000,00 EUR für Bestandsarbeiten und 3.000,00 EUR pauschalen Nebenkosten ergibt sich bei vollständigem Abruf ein Nettohonorar von 54.000,00 EUR. Gesetzliche Umsatzsteuer kommt hinzu. Änderungen der anrechenbaren Kosten allein ändern diese Pauschale nicht.",
      ("h","3 Abstimmung und Änderungen"),"Frau Döring entscheidet für den Bauherrn über Raumprogramm und Bemusterung. Vergaben und vergütungswirksame Leistungsänderungen bedürfen der gemeinsamen Bestätigung von Frau Döring und Herrn Feldmann. Die Architektin hat keine Vollmacht, Bauaufträge oder Nachträge namens des Vereins abzuschließen. Sie informiert den Vorstand vor weiterer Planung über absehbare Überschreitungen des verfügbaren Gesamtrahmens von 430.000,00 EUR brutto. Ein Erfolg unabhängig von Marktpreisen oder unerkanntem Bestand wird nicht zugesagt.",
      ("h","4 Rechnungen und Objektbetreuung"),"Abschläge dürfen nach dokumentiertem Leistungsfortschritt verlangt werden. Die Teilrechnungen bezeichnen den Bearbeitungsstand und bereits gezahlte Beträge. Nach Abschluss der Bauausführung werden die Planungs- und Überwachungsleistungen gesondert zur Abnahme vorgestellt. Die Übergabe des Gebäudes an den Hauswart ersetzt diese Erklärung nicht.",
      "In der Objektbetreuung werden gemeldete Mängel fachlich bewertet und die Verantwortlichen zur gemeinsamen Besichtigung eingeladen. Die Architektin führt eine Begehung während der laufenden Gewährleistungszeiten sowie rechtzeitig vor deren jeweiligem Ablauf durch und wirkt bei der Freigabe vereinbarter Sicherheiten mit. Die Fristen werden gewerkeweise anhand der tatsächlichen Abnahmen und Vertragsunterlagen geführt. Eine rechtliche Prüfung von Hemmung, Neubeginn oder streitiger Anspruchsinhaberschaft ist nicht als Rechtsberatung beauftragt.",
      ("h","5 Unterlagen und Vertragsende"),"Planstände werden mit Datum und Revision geführt. Bauherr und ausführende Firmen erhalten die freigegebenen Unterlagen. Frau Hartwig übergibt bei Projektabschluss bearbeitbare Plandateien und einen geordneten Dokumentennachweis. Bei Beendigung des Vertrags werden bis dahin erbrachte Leistungen nachvollziehbar abgerechnet. Die Parteien vereinbaren deutsches Recht; Allgemeine Geschäftsbedingungen werden nicht einbezogen."
    ],"Lena Hartwig, Architektin; Maren Döring und Uwe Feldmann, Vorstand",[1])
    record(3,"plan","verein","Grundlagengespräch und Ortsbesichtigung",[
      "am 04.02.2021 haben wir das Gebäude zwischen 09:00 und 11:40 Uhr mit Frau Döring, Herrn Winter und Herrn Mertens begangen. Die Räume waren wegen ausgefallener Veranstaltungen nur zeitweise beheizt. Das Protokoll hält die erhobenen Angaben und die mit dem Vorstand abgestimmten nächsten Schritte fest.",
      ("h","1 Vorgefundene Bedingungen"),"Die Außenmaße des Bestands betragen ungefähr 18,00 m mal 12,00 m. Ein Dachraum ist nur über eine Wartungsluke erreichbar und wird nicht für Veranstaltungen genutzt. Der nördliche Sockel weist im Saal einen hellgrauen Rand bis etwa 25 cm über dem Fußboden auf. Mit einem kapazitiven Messgerät wurden erhöhte relative Anzeigewerte festgestellt; daraus lässt sich kein Wassergehalt in Masseprozent ableiten. Eine belastbare Aussage zur Ursache ist ohne weitere Untersuchung nicht möglich.",
      "Herr Winter bezeichnet die im Hof sichtbare alte Leitung als stillgelegte Pumpenleitung. In der vom Verein vorgelegten Handskizze von 1992 endet sie vor der Nordwand. Ein Bestandsplan des Leitungsbetreibers ist nicht vorhanden. Frau Döring bittet, im ersten Schritt nur die beiden Anschlussstellen zu öffnen, weil der Hof noch als Stellfläche genutzt wird.",
      ("h","2 Untersuchungen und Beteiligte"),"Vereinbart werden das Bestandsaufmaß und zwei Suchschlitze unmittelbar an der Nordwand. Eine flächige Suche unter dem befestigten Hof wird zurückgestellt. Für die Tragwerksplanung soll Büro Steinbach aus Northeim beauftragt werden; Herr Mertens übernimmt die Planung für Heizung, Sanitär und Elektrokoordination. Die Beiträge sind vor Abschluss des Entwurfs zusammenzuführen. Für die Feuchtesituation wird nach Freilegung des Sockels entschieden, ob ein Baustofflabor nötig ist.",
      ("h","3 Randbedingungen"),"Die Planung geht von höchstens 120 gleichzeitig anwesenden Personen im gesamten Haus aus. Der Vereinsraum bleibt ein Besprechungsraum. Der Bauherr will das Grundstück nicht zusätzlich teilen. Öffentliche Verkehrsflächen werden durch die Rampe nicht berührt. Für die Bauaufsicht sind Nutzung, Stellplätze, Rettungswege und Anbaukubatur vorzubereiten. Die Genehmigungsfähigkeit ist heute weder beantragt noch bestätigt.",
      "Frau Hartwig legt bis 18.03.2021 zwei räumlich ausgearbeitete Varianten vor. Der Vorstand entscheidet anschließend über die Weiterplanung. Die Entscheidung über eine etwaige vollständige Sockelsanierung bleibt bis zu den Befunden offen. Herr Winter erhält eine Kopie dieses Protokolls und hält den nördlichen Wandbereich bis zum nächsten Termin frei."
    ],attachments=[1])
    record(5,"plan","verein","Vorplanung mit Varianten und Kostenschätzung",[
      "auf Grundlage der Besichtigung und der Bestandsaufnahme legen wir zwei Varianten mit gleichem Nutzungsprogramm vor. Beide erhalten den 120 m² großen Saal und den 24 m² großen Vereinsraum. Für 120 Personen werden ein ebener Zugang, zeitgemäße Sanitärbereiche und eine getrennte Technikfläche vorgesehen. Die Beträge sind Preisansätze vom März 2021, keine Unternehmerangebote.",
      ("h","1 Variante 1 Anbau nach Norden"),"Der Anbau wird 10,00 m breit und 8,00 m tief. Er enthält Foyer 32 m², Sanitär 20 m², Lager 8 m² und Technik 8 m². Die 12 m² Differenz zur Bruttofläche entfallen auf Wände. Der Hofzugang erhält zwei Rampenläufe mit zusammen 12,00 m Länge. Die vorhandene nördliche Saalwand bleibt als Innenwand erhalten; die Verbindung wird durch eine neue Öffnung hergestellt. Die Bauarbeiten greifen in den ungesicherten Leitungsbereich ein.",
      ("h","2 Variante 2 Umbau mit kleinerem Westanbau"),"Ein 6,00 m mal 8,00 m großer Westanbau nimmt die Sanitärbereiche auf. Das Foyer wird im Bestand geschaffen und verkleinert den Saal um 24 m² auf 96 m². Um das gleiche Veranstaltungsprogramm zu ermöglichen, müssen Tische ausgelagert und zwei Termine des Seniorenkreises getrennt durchgeführt werden. Der längere Weg vom Parkplatz erfordert 24 m² Rampen- und Podestfläche. Der westliche Apfelbaum entfällt.",
      ("t",[["Ansatz netto EUR","Variante 1","Variante 2"],["Baukonstruktion","160.000,00","142.000,00"],["Technische Anlagen","70.000,00","74.000,00"],["Außenanlagen","20.000,00","29.000,00"],["Ausstattung","12.000,00","15.000,00"],["Planung, Untersuchungen, Gebühren","60.000,00","58.000,00"],["Summe netto","322.000,00","318.000,00"],["Umsatzsteuer 19 Prozent","61.180,00","60.420,00"],["Summe brutto","383.180,00","378.420,00"],["Abstand zum Finanzrahmen 430.000 EUR","46.820,00","51.580,00"]]),
      ("h","3 Vergleich und Terminannahmen"),"Der Mehrbetrag der Variante 1 beträgt 4.760,00 EUR brutto. Er erhält die Saalfläche und verkürzt interne Wege. Das Risiko unbekannter Leitungen ist dort größer. In beiden Varianten ist lediglich eine örtliche Sockelabdichtung angesetzt, keine flächige nachträgliche Horizontalsperre. Die Kostenschätzung enthält keine Vorsorge für einen gegebenenfalls notwendigen Austausch des alten Kanalanschlusses.",
      "Vorgesehen sind Entwurf bis Juni 2021, Abstimmung mit der Bauaufsicht und Antrag Anfang 2022, Ausführungsplanung bis Sommer 2022, Angebote im Frühjahr 2023 und Bau ab August 2023 bis März 2024. Dieser Ablauf berücksichtigt, dass der Verein für 2022 bereits Veranstaltungen angenommen hat. Er setzt rechtzeitige Freigaben voraus. Die Architektin regt vor dem Fundamentaushub eine Leitungsortung im Anbaubereich an.",
      "Bitte teilen Sie mit, welche Variante weiterbearbeitet werden soll. Eine Entscheidung über den Verzicht auf zusätzliche Bestandserkundung wird mit der Variantenwahl nicht getroffen. Beide Varianten sind mit einer kleinen Grundrissskizze im Bestandsblatt räumlich verortet; das weiterzuentwickelnde Raumprogramm ist vorstehend vollständig aufgeführt."
    ],attachments=[4])
    record(6,"verein","plan","Weiterplanung des Nordanbaus",[
      "der Vorstand hat gestern die Varianten vom 18.03.2021 einschließlich Ihrer Kostentabelle besprochen. Wir wählen den Nordanbau mit 80 m² Bruttofläche. Für den Seniorenkreis wäre ein kleinerer Saal keine brauchbare Dauerlösung. Die Kostenschätzung von 383.180,00 EUR brutto liegt noch innerhalb unseres Rahmens von 430.000,00 EUR.",
      "Bitte arbeiten Sie den Entwurf auf dieser Grundlage aus. Im Vereinsraum bleibt es bei Besprechungen mit höchstens 18 Personen. Die beiden Suchschlitze an der Nordwand sind freigegeben. Für eine Untersuchung des gesamten Hofs möchten wir zunächst ein konkretes Angebot mit Wiederherstellung der Pflasterfläche sehen. Herr Winter meint weiterhin, dass die Pumpenleitung stillliegt; wir möchten das aber nicht als gesicherte Bestandsangabe behandeln.",
      "Die Veranstaltung im Oktober 2022 muss im Haus stattfinden können. Der Baubeginn im August 2023 passt deshalb. Die Terminabstimmung mit den Fachplanern kann jetzt erfolgen. Im Anhang übersenden wir Ihre Variantenuntersuchung unverändert als Bezug für unseren Beschluss. Diese Nachricht ruft noch keine Bauleistungen und noch nicht die weiteren Planungsphasen ab."
    ],attachments=[5])
    record(8,"plan","verein","Objektbeschreibung zum Entwurf",[
      "der Entwurf E-01 vom 10.06.2021 konkretisiert den gewählten Nordanbau. Die Bruttofläche des Bestands beträgt 216 m², der Anbau ergänzt 80 m². Die gemeinsame Bruttofläche beträgt 296 m². Die Außenabmessungen und die Baukörperlage sind im Entwurfsplan bemaßt. Die Flächen nachstehend sind Nutzungsansätze, keine vermietbaren Flächen.",
      ("h","1 Raumordnung"),"Der Saal bleibt im südlichen Bestand. Vereinsraum, Küche und Lager liegen östlich davon. Der neue Eingang führt über die Hoframpe in das 32 m² große Foyer im Nordanbau. Der Sanitärbereich mit 20 m² enthält ein barrierefrei nutzbares WC sowie weitere getrennte WC-Räume. Technik und Lager beanspruchen je 8 m². Die 68 m² nutzbare Anbaufläche und 12 m² Wandfläche ergeben zusammen die Bruttofläche.",
      ("h","2 Konstruktion und Ausbau"),"Der Anbau erhält Streifenfundamente und eine gedämmte Bodenplatte, tragende Mauerwerkswände und ein leichtes Flachdach. Der Bestand behält sein Satteldach. Eine Bewegungsfuge trennt die Baukörper. Die neue Saalöffnung erhält einen statisch bemessenen Sturz. Die Ausführung darf erst nach abgestimmtem Tragwerksplan erfolgen.",
      "Innen werden lose Sockelputze im Bestand entfernt und mineralisch erneuert. Die bisherige Feuchteerscheinung wird nicht allein durch Überstreichen behandelt. Die Kostenberechnung erfasst eine örtliche Außenabdichtung, aber keine umfassende Sanierung der historischen Bodenplatte. Die Daten aus den Suchschlitzen lassen die vorhandene Horizontalsperre weiterhin nicht sicher erkennen.",
      ("h","3 Höhen und Betrieb"),"Bezugshöhe ist der Fertigfußboden des Saals mit 112,40 m im projektspezifischen Höhensystem. Der Hof liegt am Rampenantritt bei 111,80 m. Zwei Läufe von je 6,00 m überwinden je 0,30 m. Die Eingangsschwelle wird ohne Absatz geplant. Eine unmittelbar vorgelagerte Rinne führt Niederschlagswasser zur Regenwasserleitung. Die Lage der Rinne und der seitliche Anschluss an den Altbausockel sind in der Ausführungsplanung zu detaillieren.",
      "Die Planung setzt weiterhin höchstens 120 gleichzeitig anwesende Personen und Besprechungsnutzung im Vereinsraum voraus. Eine spätere Nutzung als Gymnastik-, Tanz- oder Veranstaltungsraum mit anderer Personenzahl ist nicht Bestandteil dieser Freigabe. Die neue Technik wird so bemessen, dass die festgelegten Betriebszeiten zwischen 08:00 und 22:00 Uhr abgedeckt sind."
    ],attachments=[7])
    record(10,"plan","verein","Abstimmung mit den Fachplanern",[
      "am 15.06.2021 wurden der Entwurf und die Kostenberechnung mit Paul Mertens, Tragwerksplaner Dirk Steinbach und dem Vorstand abgestimmt. Die Besprechung fand von 10:00 bis 12:15 Uhr im Bürgerhaus statt. Die nachstehenden Festlegungen werden in die Bauvorlagen übernommen.",
      ("h","1 Tragwerk und Gebäudetechnik"),"Herr Steinbach legt für den Nordanbau eine vom Bestand getrennte Gründung zugrunde. Die Saalöffnung wird mit einem Stahlträger überdeckt; dessen Auflager dürfen nicht im geschädigten Sockelputz bemessen werden. Die Tragwerksunterlagen T-02 werden vor Freigabe der Ausführung durch die Architektin mit den Öffnungen im Gebäudeplan abgeglichen. Es werden keine Lasten des Anbaus in das alte Dach eingeleitet.",
      "Herr Mertens sieht die neue Verteilung im nordöstlichen Technikraum vor. Abwasser quert die Fundamente nur in festgelegten Hülsen. Der Altstrang im Hof soll vor Baubeginn geortet werden. Der Verein beauftragt heute noch keine flächige Untersuchung; Frau Döring will den Umfang mit dem Hauswart besprechen. Dieser Punkt bleibt mit Verantwortlichem Bauherr und Wiedervorlage zur Ausführungsfreigabe offen.",
      ("h","2 Kostenabgleich"),"Die Kostenberechnung beträgt 337.240,00 EUR netto und 401.315,60 EUR brutto. Gegenüber der Schätzung vom 18.03.2021 ergibt sich ein Anstieg von 15.240,00 EUR netto beziehungsweise 18.135,60 EUR brutto. Hauptursachen sind konkretisierte Gründung, 74.000,00 EUR technische Anlagen und die Sockelarbeiten. Zum Finanzrahmen verbleiben 28.684,40 EUR brutto. Der Vorstand bestätigt den Entwurf, ohne den Gesamtrahmen anzuheben.",
      ("h","3 Genehmigung und Ablauf"),"Für den Bauantrag werden Lageplan, Grundriss, Ansichten, Schnitt und Baubeschreibung bis Januar 2022 zusammengestellt. Das Nutzungskonzept bleibt unverändert. Rettungswege aus dem Saal führen getrennt nach Süden und über das Foyer nach Norden. Die Verbindungstür darf nicht als Lagerfläche genutzt werden. Herr Mertens übergibt die Angaben zum Technikraum bis Ende Juni; die Architektin übernimmt sie in die Antragsfassung.",
      "Der Vorstand erhält Kostenberechnung und Entwurfsplan. Die endgültige Bestellung von Bauleistungen erfolgt nach Genehmigung und Angebotsprüfung. Die Beteiligten stimmen dem Protokollinhalt zu; Herr Winter erhält den Abschnitt zum Leitungsbestand zur Kenntnis."
    ],"Lena Hartwig; Paul Mertens; Dirk Steinbach; Maren Döring",[7,9])
    record(11,"verein","amt","Bauantrag Umbau und Nordanbau",[
      "wir beantragen die Genehmigung des Umbaus und der Erweiterung unseres Bürgerhauses in Einbeck. Der Antrag bezieht sich auf die nachstehend eindeutig bezeichnete Baumaßnahme und die beigefügten Bauvorlagen. Zustellungsbevollmächtigte für Rückfragen und Bescheide ist Lena Hartwig, Mühlenweg 7, 37574 Einbeck.",
      ("h","1 Antragsteller und Baugrundstück"),("t",[["Angabe","Eintragung"],["Bauherr und Eigentümer","Bürgerhaus Leinewinkel e.V., Lindengasse 18, 37574 Einbeck"],["Vertretung","Maren Döring und Uwe Feldmann, Vorstand"],["Grundstück","Gemarkung Einbeck, Flur 14, Flurstück 83/7"],["Grundstücksfläche","1.512 m²; 42,00 m mal 36,00 m"],["Entwurfsverfasserin","Lena Hartwig, Architektin, Mühlenweg 7, 37574 Einbeck"],["Bauvorhaben","Umbau eines eingeschossigen Bürgerhauses und 80 m² Nordanbau"]],[145,350]),
      ("h","2 Umfang des Antrags"),"Beantragt werden eine neue Foyer- und Sanitärzone, der niveaugleiche Hauptzugang über eine Rampe, der Durchbruch zwischen Saal und Foyer sowie die Erneuerung der Innenoberflächen und technischen Anlagen. Der Saal wird nicht vergrößert. Im Haus werden höchstens 120 Personen gleichzeitig aufgenommen. Die Nutzung umfasst Vereinsbesprechungen, Chorproben, Seniorennachmittage und gelegentliche private Feiern ohne Übernachtung.",
      "Der Bestand hat 216 m², der Anbau 80 m² Bruttofläche. Die überbaute Grundfläche beträgt nach Ausführung 296 m². Die Grundstückszufahrt erfolgt unverändert von der Lindengasse. Auf dem eigenen Grundstück werden zwölf Kfz-Stellplätze einschließlich eines verbreiterten Stellplatzes und acht Fahrradplätze nachgewiesen. Der Anbau bleibt 8,00 m von der nördlichen Grundstücksgrenze entfernt.",
      ("h","3 Erklärungen und Bauvorlagen"),"Die im Lageplan eingetragenen Grundstücksgrenzen und Maße entsprechen den vom Bauherrn zur Planung vorgelegten Unterlagen. Änderungen des Nutzungsprogramms werden vor ihrer Ausführung nachgereicht. Baubeginn ist für August 2023 vorgesehen. Die Auswahl ausführender Unternehmen ist noch nicht erfolgt. Der Bauherr bestätigt die Verfügungsbefugnis über das gesamte Baugrundstück.",
      "Die Unterzeichnenden bestätigen die Übereinstimmung von Lageplan L-01, Bauzeichnung G-01 und Baubeschreibung B-01. Die Angaben zum Brandschutz sind in der Baubeschreibung enthalten. Tragwerksplanung T-02 und technische Ausführungsnachweise werden vor Beginn der betreffenden Arbeiten vorgelegt. Die Aktenführung soll unter der Projektbezeichnung BH-21-04 erfolgen."
    ],"Maren Döring und Uwe Feldmann, Bauherr; Lena Hartwig, Entwurfsverfasserin",[12,13,14])
    record(14,"plan","amt","Baubeschreibung zum Bauantrag",[
      "die Baubeschreibung B-01 gehört zum Antrag vom 20.01.2022 für das Bürgerhaus Lindengasse 18, Gemarkung Einbeck, Flur 14, Flurstück 83/7. Sie beschreibt das gesamte beantragte Vorhaben und die künftige Betriebsweise, nicht nur den neuen Gebäudeteil.",
      ("h","1 Gebäude und Nutzung"),"Das eingeschossige Gebäude wird ohne Keller genutzt. Die Oberkante Fertigfußboden liegt 0,60 m über dem Hofzugang. Der Bestand misst 18,00 m mal 12,00 m und hat einen nicht ausgebauten Dachraum. Der nördliche Anbau misst 10,00 m mal 8,00 m und erhält ein Flachdach mit Attika bei 3,50 m über dem Saalfußboden. Es entstehen keine Wohnungen und keine Übernachtungsräume.",
      "Für die Nutzung werden maximal 120 Personen einschließlich Personal festgelegt. Der Saal umfasst 120 m², der Vereinsraum 24 m². Der neue Foyerbereich umfasst 32 m², Sanitär 20 m² sowie Lager und Technik jeweils 8 m². Veranstaltungen enden um 22:00 Uhr. Ein regelmäßiger Tanz-, Fitness- oder Gaststättenbetrieb ist nicht beantragt; Speisen werden nur angeliefert oder in der bestehenden Teeküche angerichtet.",
      ("h","2 Baustoffe und Erschließung"),"Der Anbau erhält Stahlbetonfundamente, eine gedämmte Bodenplatte, mineralisches Mauerwerk und ein leichtes gedämmtes Dach. Zwischen Bestand und Anbau wird eine Bewegungsfuge angeordnet. Das vorhandene Satteldach bleibt bestehen. Die tragende Öffnung zum Foyer wird nach gesondertem Tragwerksnachweis hergestellt. Dachwasser wird über einen getrennten Regenwasserstrang geführt. Schmutzwasser wird an die bestehende Grundstücksentwässerung angeschlossen.",
      ("h","3 Zugang und Rettungswege"),"Vom verbreiterten Stellplatz führt ein 1,50 m breiter Weg zur Rampe. Zwei Rampenläufe von je 6,00 m mit Zwischenpodest überwinden 0,60 m Höhe. Beidseitige Handläufe und Radabweiser werden ausgeführt. Die drei Außentüren haben jeweils 1,00 m lichte Durchgangsbreite. Der Saal besitzt zwei voneinander getrennte Ausgänge: direkt nach Süden und über das Foyer nach Norden. Türen in Fluchtrichtung sind während der Nutzung ohne Schlüssel zu öffnen.",
      "Die Sanitärzone enthält einen von beiden Nutzungsgruppen unabhängig erreichbaren barrierefrei nutzbaren WC-Raum. Der Bewegungsbereich vor den Türen bleibt frei. Es werden zwölf Kfz-Stellplätze auf dem Grundstück vorgehalten; acht Fahrradplätze liegen westlich der Rampe. Der nordseitige Rettungsweg endet auf dem eigenen Hof und erreicht über den westlichen Weg die öffentliche Straße.",
      ("h","4 Bauvorbereitung"),"Vor Erdarbeiten wird der Leitungsbestand im Arbeitsbereich festgestellt. Die bisher nur punktuell untersuchte Bestandswand wird abschnittsweise freigelegt. Feuchte- und Salzbefunde sind vor dem Neuverputz zu dokumentieren. Der Bauherr nutzt das Gebäude während des Umbaus nicht für Veranstaltungen. Die Bauleitung stimmt die erforderlichen Anzeigen und Nachweisvorlagen mit der Bauaufsicht ab."
    ],"Lena Hartwig, Entwurfsverfasserin; Maren Döring, Bauherrin",[12,13])
    record(15,"amt","plan","Nachforderung zum Bauantrag 63-2022-041",[
      "der Bauantrag für das Bürgerhaus Lindengasse 18 vom 20.01.2022 ist unter dem Aktenzeichen 63-2022-041 erfasst. Für die weitere Bearbeitung benötigen wir eine ergänzte Darstellung der Rettungswege aus dem Vereinsraum und der Begrenzung der Gesamtbelegung. Die Bezeichnung Besprechungsraum allein genügt dafür nicht.",
      "Bitte geben Sie die lichte Breite der Türen, die Lauflänge bis ins Freie und die Betriebsweise bei gleichzeitig genutztem Saal und Vereinsraum an. Der Bauherr soll bestätigen, wie die höchstens 120 anwesenden Personen einschließlich Helfern erfasst werden. Die Rampe muss auch während Veranstaltungen frei bleiben. Im bisherigen Grundriss ist der nördliche Ausgang dargestellt; die Möblierung vor diesem Ausgang ist nicht eingetragen.",
      "Reichen Sie außerdem eine Zuordnung der Raumbezeichnungen zwischen Grundriss und Baubeschreibung nach. Änderungen der Kubatur sind nach heutigem Bearbeitungsstand nicht verlangt. Bitte verwenden Sie für die Ergänzung die vorhandene Projektnummer, damit die Unterlagen dem Antrag zugeordnet werden können. Vor Erteilung der Genehmigung darf nicht mit der beantragten Baumaßnahme begonnen werden. Es sind keine Dateien an diese Nachricht angehängt."
    ])
    record(16,"plan","amt","Ergänzung zum Rettungswegkonzept",[
      "zu Ihrer Nachforderung vom 03.02.2022 reichen wir die Ergänzung BS-02 zum Antrag 63-2022-041 ein. Die Baukörperlage und die Flächen des Bauantrags bleiben unverändert. Der Bauherr hat die nachstehende Belegungs- und Betriebsorganisation bestätigt.",
      ("h","1 Belegung und Raumzuordnung"),"Im Saal S-01 werden bei Bestuhlung 96 Gäste zugelassen. Bis zu sechs Helfer halten sich zusätzlich dort auf. Der Vereinsraum V-01 wird mit höchstens 18 Personen genutzt. Bei dieser gleichzeitigen Nutzung ist die Gesamtgrenze von 120 Personen ausgeschöpft. Der Foyerbereich ist Durchgangsfläche und wird nicht zusätzlich vermietet. Bei anderer Belegung verringert der Hauswart die zugelassene Personenzahl im Saal entsprechend.",
      ("h","2 Wege und Türen"),("t",[["Raum","Erster Weg","Weiterer Weg"],["Saal S-01","Südausgang T-01, 1,00 m licht, längster Weg 16,50 m","Durch T-02 in Foyer, Nordausgang T-03, jeweils 1,00 m licht, insgesamt 22,00 m"],["Vereinsraum V-01","Über Bestandsflur zum Südausgang, längster Weg 19,00 m","Über Verbindung zum Saal und Foyer zum Nordausgang, 24,00 m"],["Foyer F-01","Nordausgang T-03 direkt ins Freie","Durch T-02 und Saal zum Südausgang"]],[85,205,205]),
      "Die Türen T-01 und T-03 öffnen nach außen. T-02 erhält einen während der Betriebszeit ohne Schlüssel bedienbaren Beschlag. Der lichte Weg von 1,20 m durch den Bestandsflur und der freie Bereich vor T-03 dürfen nicht durch Garderoben oder abgestellte Tische eingeengt werden. Der Möblierungsplan wird vom Hauswart beim Aufbau der Veranstaltungen kontrolliert.",
      ("page",), ("h","3 Betrieb und technische Ausstattung"),"Der Verein benennt Ralf Winter als verantwortliche Person für Öffnen, Belegungskontrolle und Schlusskontrolle. Bei jeder Vermietung wird die Personenzahl schriftlich eingetragen. Die Küche bleibt Teeküche ohne Fritteuse, offenen Grill oder Gasgeräte. Feuerlöscher werden an den beiden Ausgängen angeordnet; die Kennzeichnung und Sicherheitsbeleuchtung der Wege wird mit der Elektroplanung umgesetzt.",
      "Der Rettungsweg über den Hof ist auch während Winterveranstaltungen freizuhalten. Die laufende Unterhaltung des Wegs übernimmt der Verein. Eine Nutzung des Vereinsraums als regelmäßiger Gymnastikraum ist nicht Grundlage dieser Ergänzung. Die Architektin erklärt, dass die vorstehenden Tür- und Wegangaben mit der Bauzeichnung G-01 übereinstimmen. Der Bauherr verpflichtet sich zur beschriebenen Betriebsorganisation."
    ],"Lena Hartwig, Architektin; Maren Döring, Vorsitzende",[13,14])
    record(17,"amt","verein","Baugenehmigung 63-2022-041",[
      "auf Ihren Antrag vom 20.01.2022 wird die Genehmigung für den Umbau und den nördlichen Anbau des Bürgerhauses auf dem Grundstück Lindengasse 18, Gemarkung Einbeck, Flur 14, Flurstück 83/7, nach Paragraf 70 NBauO erteilt. Maßgeblich sind die nachfolgend bezeichneten Bauvorlagen und die Nebenbestimmungen dieses Bescheids.",
      ("h","1 Genehmigter Umfang"),"Genehmigt ist die Erweiterung des eingeschossigen Bestands von 216 m² um einen 80 m² großen Anbau mit Foyer, Sanitär-, Lager- und Technikräumen. Die Gesamtbelegung wird auf 120 Personen einschließlich Helfern begrenzt. Der Vereinsraum wird als Besprechungsraum mit höchstens 18 Personen genutzt. Ein regelmäßiger Tanz-, Fitness- oder Gaststättenbetrieb und Übernachtungen sind nicht Gegenstand dieser Genehmigung.",
      ("h","2 Zugehörige Bauvorlagen"),("t",[["Kennung","Unterlage","Stand"],["L-01","Lageplan mit Stellplätzen und Grundstücksgrenzen","20.01.2022"],["G-01","Grundriss, Schnitt und Ansichten","20.01.2022"],["B-01","Baubeschreibung","20.01.2022"],["BS-02","Ergänzung zum Rettungswegkonzept","14.02.2022"]]),
      ("h","3 Nebenbestimmungen"),"Vor Beginn der tragenden Arbeiten ist der Standsicherheitsnachweis für Gründung, Saalöffnung und Dachkonstruktion bei der Bauaufsicht vorzulegen. Abweichungen von den eingereichten Öffnungsmaßen oder der Baukörperlage sind vor ihrer Ausführung mitzuteilen. Mit den betroffenen Arbeiten darf erst begonnen werden, wenn die erforderlichen Nachweisvorlagen erfolgt sind.",
      "Die Rettungswege und Türbreiten aus BS-02 sind während des Betriebs dauerhaft freizuhalten. Die Außentüren müssen von innen ohne Schlüssel geöffnet werden können. Die stufenfreie Erreichbarkeit des Foyers und des hierfür vorgesehenen WC-Raums ist vor Aufnahme der Nutzung herzustellen. Der Betrieb darf erst aufgenommen werden, wenn diese Einrichtungen und die erforderliche Rettungswegkennzeichnung vorhanden sind.",
      "Die Belegungsgrenze gilt für das gesamte Haus. Der Hauswart hat die bei Vermietung vereinbarte Personenzahl zu dokumentieren. Eine Nutzungsänderung oder eine andere Belegungsorganisation ist vor Umsetzung mit der Bauaufsicht abzustimmen. Die zwölf im Lageplan bezeichneten Stellplätze und der Zugang von der Straße sind auf dem Grundstück zu erhalten.",
      ("h","4 Hinweise und Kosten"),"Private Rechte Dritter bleiben unberührt. Die Genehmigung ersetzt keine privatrechtliche Zustimmung für Eingriffe in fremde Leitungen. Die Kostenentscheidung ergeht gesondert. Die Bauaufsicht erhält die Benennung der Bauleitung und die Anzeige des vorgesehenen Baubeginns. Die genehmigten Unterlagen müssen auf der Baustelle zugänglich sein.",
      ("h","5 Rechtsbehelfsbelehrung"),"Gegen diesen Bescheid kann innerhalb eines Monats nach Bekanntgabe Widerspruch bei der Stadt Einbeck, Teichenweg 1, 37574 Einbeck, erhoben werden. Der Widerspruch ist schriftlich, in der gesetzlich zugelassenen elektronischen Form oder zur Niederschrift einzulegen. Eine einfache E-Mail genügt der vorgeschriebenen Form nicht."
    ],"Klara Brandes, im Auftrag",[12,13,14,16])
    record(18,"verein","plan","Abruf weiterer Planung und Überwachung",[
      "nach Eingang der Baugenehmigung vom 17.03.2022 rufen wir die im Vertrag vom 18.01.2021 vereinbarten Leistungsphasen 5 bis 9 ab. Die Vergütung richtet sich unverändert nach diesem Vertrag. Herr Feldmann stimmt diesem Abruf als zweites vertretungsberechtigtes Vorstandsmitglied zu. Der verfügbare Gesamtrahmen bleibt bei 430.000,00 EUR brutto.",
      "Bitte führen Sie die Ausführungsplanung einschließlich Rampenanschluss, Rinne und Sockelübergang fort. Den Rahmenterminplan vom 25.04.2022 bestätigen wir als Abstimmungsgrundlage. Die Veranstaltung im Oktober 2022 bleibt unangetastet. Die Angebote sollen erst im Frühjahr 2023 eingeholt werden; die aktuelle Kostenberechnung ersetzt dafür keine Marktpreise.",
      "Die Tragwerksunterlagen müssen vor den tragenden Arbeiten bei der Bauaufsicht vorliegen. Bitte erinnern Sie uns vor Ausschreibung nochmals an die Leitungssuche. Wir haben bisher nur die beiden Suchschlitze bezahlt, keine flächige Ortung. Das Haus wird ab dem 21.08.2023 vollständig für die Bauarbeiten geräumt. Im Anhang befinden sich die Genehmigung und der bestätigte Rahmenterminplan. Die Vergabevollmacht verbleibt beim Vorstand."
    ],attachments=[17,19])

    lv_intro=[
      "für das Bürgerhaus Lindengasse 18 in Einbeck wird Los 1 Rohbau und Gebäudehülle zur Angebotsabgabe versandt. Auftraggeber ist der Bürgerhaus Leinewinkel e.V. Dieses Leistungsverzeichnis ist ein vollständiges Los; Haustechnik, Fenster, Innentüren, Bodenbeläge und lose Ausstattung werden gesondert vergeben. Der Bauherr ist ein privat finanzierter Verein. Es handelt sich um eine private Angebotsanfrage.",
      ("h","1 Vertrags- und Ausführungsbedingungen"),
      "Die Ausführung beginnt am 21.08.2023. Die vertraglich vorgesehene Fertigstellung ist der 15.03.2024. Abgerechnet werden die tatsächlich nachgewiesenen Mengen zu den angebotenen Einheitspreisen. Die VOB wird nicht als Vertragsbestandteil vereinbart. Stoffe und Nebenleistungen, die in der jeweiligen Position ausdrücklich genannt sind, sind im Einheitspreis enthalten. Pauschale Nebenangebote ohne Abgrenzung werden nicht gewertet.",
      "Ausführungsgrundlagen sind AP-01 vom 02.06.2022 und D-12 vom 06.06.2022. Die technische Bearbeitung der Haustechnik vom 07.03.2023 ändert diese freigegebenen Gebäudepläne nicht. Vor verdeckenden Arbeiten ist die Bauleitung zur Kontrolle einzuladen. Der Verein beauftragt eine Leitungsortung vor Aushub; der Auftragnehmer prüft erkennbare Leitungen trotzdem und hält bei unklaren Funden Rücksprache. Aus der fehlenden Eintragung im Plan folgt keine Freigabe zum Durchtrennen.",
      "Die Angebote sind bis 04.04.2023 um 12:00 Uhr an den Vorstand zu richten. Preise sind in EUR netto, Umsatzsteuer gesondert anzugeben. Die Bindefrist endet am 28.04.2023. Mit dem Angebot sind Ausführbarkeit, Personalansatz und Vorlauf für Türen und Dachmaterial zu bestätigen. Gerüste für die eigenen Leistungen und die Entsorgung eigener Verpackungen sind enthalten.",
      ("h","2 Leistungsverzeichnis")]
    for i,(name,desc) in enumerate(ITEMS):
        lv_intro += [("h",f"2.{i+1} {name}"),desc,
                     f"Menge: {str(QTY[i]).replace('.',',')} {UNITS[i]}. Einheitspreis und Gesamtpreis sind vom Bieter in seinem Angebot anzugeben."]
    lv_intro += [("h","3 Mengenermittlung und Schnittstellen"),
      "Die Rückbaufläche folgt dem Bestandsaufmaß, die Bodenplatte der Anbaugrundfläche. Beim Mauerwerk wurden Öffnungen abgezogen. Die Abdichtungsfläche setzt 36 m² freigelegten Sockel voraus. Der Innenputzansatz erfasst 144 m² Wandfläche, nicht die Bodenfläche. Die Rampe umfasst 18,00 m² Lauf- und 2,25 m² Podestfläche. Zusätzliche Anschlüsse an die Hofpflasterung werden gesondert aufgemessen.",
      "Elektro und Sanitär liefern die erforderlichen Hülsenmaße vor Betonage. Für die Dichtheit der von Los 1 gelieferten Rinnenanschlüsse ist Los 1 zuständig, für die anschließende unterirdische Grundstücksleitung das gesonderte Tiefbauunternehmen. Die Schnittstelle liegt am ersten Rohrstutzen hinter der Rinne. Über die Bauleitung hinausgehende rechtsgeschäftliche Anordnungen dürfen nur Frau Döring und Herr Feldmann gemeinsam erteilen."]
    record(23,"plan","bau","Leistungsverzeichnis Los 1",lv_intro,attachments=[20,21])
    for n,key,eps in [(25,"bau",LEINE_EP),(26,"hage",HAGE_EP)]:
        total=round(sum(q*p for q,p in zip(QTY,eps)),2)
        rows=[["Position / Leistung","Menge","Einheit","EP EUR","Gesamt EUR"]]
        rows += [[f"2.{i+1} {ITEMS[i][0]}",str(q).replace('.',','),UNITS[i],euro(eps[i]),euro(q*eps[i])] for i,q in enumerate(QTY)]
        rows += [["Summe netto","","","",euro(total)],["Umsatzsteuer 19 Prozent","","","",euro(total*.19)],["Angebot brutto","","","",euro(total*1.19)]]
        record(n,key,"verein",f"Angebot {'LB-230404' if n==25 else 'HB-23041'} für Los 1",[
          "auf Ihre Anfrage vom 08.03.2023 bieten wir das Los 1 für Umbau und Anbau des Bürgerhauses Lindengasse 18 zu den folgenden Einheitspreisen an. Die Langtexte des Leistungsverzeichnisses und die dort benannten freigegebenen Gebäudepläne sind Bestandteil unseres Angebots. Mengenänderungen werden zu den angebotenen Einheitspreisen abgerechnet.",
          ("h","1 Angebotspreise"),("t",rows,[210,55,40,80,110]),
          ("h","2 Ausführung und Bindung"),
          ("Wir bestätigen Baubeginn 21.08.2023 und Fertigstellung 15.03.2024. Für die Außentüren beträgt der heutige Beschaffungsvorlauf sechs Wochen ab Aufmaß und Freigabe. Geplant sind ein Polier und drei Fachkräfte; Beton- und Dacharbeiten werden mit eigenem Personal ausgeführt. Die Bindefrist 28.04.2023 wird bestätigt." if n==25 else
           "Wir können den vorgesehenen Baubeginn am 21.08.2023 einhalten. Für die Außentüren rechnen wir mit zwölf Wochen ab Freigabe. Die Fertigstellung bis 15.03.2024 ist bei Freigabe der Türmaße bis 15.11.2023 möglich. Geplant sind ein Vorarbeiter und vier Fachkräfte. Das Angebot bleibt bis 28.04.2023 verbindlich."),
          "Die ausgewiesenen Einheitspreise enthalten die beschriebenen Nebenleistungen, Materiallieferung und eigene Entsorgung. Ein Nachlass und Skonto werden nicht angeboten. Eine automatische Stoffpreisgleitung wird nicht verlangt. Rechnungen sind binnen 21 Kalendertagen nach prüffähigem Zugang zahlbar. Abschläge werden nur für nachgewiesene Leistungen gestellt.",
          ("h","3 Bestand und Abgrenzung"),
          ("Die Sockelposition umfasst den beschriebenen vertikalen Abdichtungsaufbau, keine Injektion einer Horizontalsperre. Die vorhandenen Leitungen müssen vor Beginn bezeichnet sein. Eine unbekannte Leitung wird nicht auf eigene Entscheidung entfernt. Der angebotene Preis enthält keine Stillstandstage aufgrund fremder Leitungsfreigaben." if n==25 else
           "Unser Angebot umfasst auch die Schutzlage und Wiederverfüllung der Sockelabdichtung. Mit dem Ansatz ist keine nachträgliche Horizontalsperre verbunden. Nicht angeboten sind Leitungsumlegungen außerhalb der im Leistungsverzeichnis beschriebenen Arbeiten. Sollte eine unbekannte Leitung angetroffen werden, werden wir deren Klärung anzeigen und den betroffenen Bereich sichern."),
          "Bitte erteilen Sie einen schriftlichen Auftrag mit Benennung der freigegebenen Planstände. Unsere Kalkulation beruht auf dem vollständigen Leistungsverzeichnis und nicht auf einer verkürzten Portalvorschau. Weitere Geschäftsbedingungen fügen wir nicht bei."
        ],signer=("Henning Seidel, Geschäftsführer" if n==25 else None),attachments=[23])
    plan=sum(q*p for q,p in zip(QTY,PLAN_EP));le=sum(q*p for q,p in zip(QTY,LEINE_EP));ha=sum(q*p for q,p in zip(QTY,HAGE_EP))
    record(27,"plan","verein","Angebotsprüfung und Vergabevorschlag",[
      "die fristgerecht eingegangenen Angebote Leinebau vom 04.04.2023 und Hagedorn vom selben Tag wurden anhand des Leistungsverzeichnisses rechnerisch, mengenbezogen und auf Leistungsabgrenzungen geprüft. Es liegen zwei wertbare Hauptangebote vor. Die Prüfung betrifft die private Vergabe des Vereins und enthält keine Aussage über ein öffentliches Vergabeverfahren.",
      ("h","1 Preisspiegel netto EUR"),("t",[["Position","Planeransatz","Leinebau","Hagedorn"]]+[[f"2.{i+1} {ITEMS[i][0]}",euro(QTY[i]*PLAN_EP[i]),euro(QTY[i]*LEINE_EP[i]),euro(QTY[i]*HAGE_EP[i])] for i in range(12)]+[["Summe",euro(plan),euro(le),euro(ha)]],[230,88,88,89]),
      ("h","2 Aufklärung am 06.04.2023"),"Herr Seidel bestätigte telefonisch und im gemeinsamen Gespräch, dass Schutzlage und Wiederverfüllung der Sockelabdichtung im Einheitspreis enthalten sind. Herr Hagedorn bestätigte die identische Leistung. Beide schließen eine Horizontalsperre aus, wie es auch der Langtext tut. Keiner hat einen Nachlass angeboten. Die Rinnenanschlussschnittstelle am ersten Rohrstutzen wurde von beiden Bietern bestätigt. Die Gesprächsteilnehmer waren Lena Hartwig, Uwe Feldmann, Henning Seidel und zeitversetzt Jens Hagedorn.",
      f"Leinebau liegt {euro(le-plan)} EUR netto über dem Planeransatz und {euro(ha-le)} EUR netto unter Hagedorn. Die Mengen und Umsatzsteuersummen sind rechnerisch richtig. Der längere Türvorlauf bei Hagedorn ist innerhalb des geplanten Ablaufs möglich, lässt aber weniger Spielraum bei späteren Maßänderungen. Eine von der Planung abweichende Ausführung wurde nicht angeboten.",
      ("h","3 Vergabevorschlag und Kostenkontrolle"),"Wir empfehlen den Auftrag an Leinebau zum geprüften Angebot. Vor Versand sind Vertragsfristen, Zahlungsfrist und Planstände ausdrücklich zu bestätigen. Für die weiteren Lose bleiben Angebote erforderlich. Das Los 1 darf nicht mit den vollständigen Baukonstruktionskosten von 169.940,00 EUR netto verwechselt werden: Fenster, Bodenbeläge, Innentüren und weitere Ausbauarbeiten liegen außerhalb dieses Loses.",
      "Die Bauleitung erhält keine selbstständige Nachtragsvollmacht. Der Auftrag soll klarstellen, dass geänderte Mengen aufgemessen werden und zusätzliche Leistungen vor Ausführung preislich beschrieben werden müssen. Die Entscheidung über den Zuschlag verbleibt beim Vorstand. Frau Döring und Herr Feldmann erhalten beide Angebote und den bepreisten Planeransatz."
    ],attachments=[24,25,26])
    record(28,"verein","bau","Bauvertrag und Zuschlag für Los 1",[
      "Bürgerhaus Leinewinkel e.V., vertreten durch Maren Döring und Uwe Feldmann, nimmt das Angebot LB-230404 der Leinebau Hochbau GmbH an. Leinebau bestätigt den Auftrag durch Henning Seidel. Für das Los 1 gelten abschließend die folgenden Vereinbarungen.",
      ("h","1 Leistung und Unterlagen"),"Geschuldet sind sämtliche zwölf Positionen des Leistungsverzeichnisses vom 08.03.2023 nach ihren vollständigen Langtexten. Vertragsgrundlagen sind dieser Vertrag, das Angebot LB-230404, das Leistungsverzeichnis sowie AP-01 vom 02.06.2022 und D-12 vom 06.06.2022 in dieser Reihenfolge. Bei technischen Widersprüchen ist vor Ausführung Rücksprache zu halten. Die VOB und sonstige Geschäftsbedingungen werden nicht einbezogen.",
      f"Die vorläufige Auftragssumme beträgt {euro(le)} EUR netto zuzüglich {euro(le*.19)} EUR Umsatzsteuer, insgesamt {euro(le*1.19)} EUR brutto. Es handelt sich um einen Einheitspreisvertrag. Abgerechnet werden die tatsächlich ausgeführten und gemeinsam festgestellten Mengen. Die Auftragssumme ist deshalb keine unveränderliche Pauschale.",
      ("h","2 Termine und Mitwirkung"),"Die Arbeiten beginnen am 21.08.2023 und sind bis 15.03.2024 fertigzustellen. Das Gebäude wird ab Baubeginn nicht für Veranstaltungen genutzt. Der Auftraggeber gewährt Zugang, benennt vorhandene Leitungen und veranlasst die im Ausführungsplan vorgesehene Ortung. Leinebau sichert erkennbare Leitungen und zeigt unklare Bestandsverhältnisse unverzüglich mit betroffenen Arbeiten, Personal und voraussichtlichen Folgen an. Eine Vertragsstrafe wird nicht vereinbart.",
      ("h","3 Änderungen und Dokumentation"),"Zusätzliche Leistungen werden dem Vorstand vor Ausführung mit Mengen, Preisen und Terminfolgen vorgelegt. Frau Döring und Herr Feldmann entscheiden gemeinsam. Frau Hartwig koordiniert technisch und prüft Angebote, darf aber den Verein nicht durch Nachtragsaufträge verpflichten. Bei akuter Gefährdung dürfen erforderliche Sicherungsmaßnahmen ausgeführt werden; Umfang und Anlass sind noch am selben Tag zu dokumentieren.",
      "Aufmaße sind vor dem Verdecken gemeinsam aufzunehmen. Tagesberichte nennen Personalstärke, Witterung, ausgeführte Arbeiten und Unterbrechungen. Bauleitungsunterschriften auf Aufmaßen bestätigen zunächst die festgestellten Maße; eine zusätzliche Vergütungsvereinbarung entsteht dadurch nicht. Fotos sind mit Ort und Datum zuzuordnen.",
      ("h","4 Vergütung und Sicherheiten"),"Prüffähige Rechnungen sind binnen 21 Kalendertagen nach Zugang zahlbar. Abschläge werden auf die Schlussrechnung angerechnet und mit Nettobetrag sowie Umsatzsteuer ausgewiesen. Es werden weder Skonto noch pauschaler Gewährleistungseinbehalt vereinbart. Bei festgestellten Mängeln bleiben die gesetzlichen Rechte unberührt. Eine Vorauszahlung oder Vertragserfüllungsbürgschaft ist nicht geschuldet.",
      ("h","5 Abnahme und Mängel"),"Nach Fertigstellung findet ein gemeinsamer Abnahmetermin statt. Der Auftraggeber erklärt die Abnahme und bezeichnet bekannte Mängel sowie Vorbehalte im Protokoll. Die Übergabe von Schlüsseln, die Rechnungsprüfung und eine bloße technische Freigabe ersetzen die vereinbarte förmliche Abnahme nicht. Für die von Leinebau hergestellten Bauleistungen wird eine fünfjährige Mängelfrist ab ihrer Abnahme vereinbart. Beginn, Änderungen oder Unterbrechungen sind anhand des jeweiligen Vorgangs zu beurteilen; Fristen anderer Unternehmer werden dadurch nicht festgelegt.",
      ("h","6 Abschluss"),"Leinebau übergibt Revisionsangaben, Produktnachweise, Bedienhinweise und die für sein Los erforderlichen Prüfprotokolle. Änderungen dieses Vertrags werden in Textform dokumentiert. Deutsches Recht gilt. Beide Parteien erhalten eine gleichlautende Ausfertigung einschließlich der im Anlagenverzeichnis benannten Unterlagen."
    ],"Maren Döring und Uwe Feldmann, Vorstand; Henning Seidel, Geschäftsführer",[20,21,23,25])
    record(29,"verein","plan","Neue Anfrage zur Nutzung des Vereinsraums",[
      "der örtliche Bewegungskreis fragt an, ob er den Vereinsraum nach Wiedereröffnung an zwei Vormittagen in der Woche für Gymnastik nutzen kann. Gedacht ist an zwölf Teilnehmerinnen und eine Übungsleiterin, ohne Musikverstärker. Im ursprünglichen Konzept stand nur Besprechungsnutzung. Eine Zusage haben wir dem Kreis noch nicht erteilt.",
      "Herr Winter hält den bisherigen Boden für zu hart und schlägt einen elastischen Belag vor. Außerdem sollen die Schränke in das neue Lager. Bitte sagen Sie uns, ob dies Änderungen an Raumakustik, Lüftung, Rettungswegen oder dem genehmigten Nutzungskonzept auslöst. Die bisherige Grenze von 120 Personen insgesamt möchten wir nicht überschreiten. Wir wünschen noch keine Ausführung zusätzlicher Leistungen.",
      "Herr Feldmann möchte zunächst eine belastbare Kostenangabe und die Rückmeldung der Bauaufsicht. Der Bauantrag und die Ergänzung BS-02 sind dafür maßgeblich. In der nächsten Baubesprechung soll geklärt werden, ob der geplante Boden auch bei unveränderter Besprechungsnutzung genügt. Die neue Nutzung darf nicht als bereits freigegeben in die Bestellung der Ausbauarbeiten eingehen. Bitte führen Sie diesen Wunsch getrennt von den jetzt laufenden Rohbauarbeiten."
    ])
    record(31,"bau","verein","Behinderungsanzeige wegen Leitungsfund",[
      "am 05.09.2023 um 08:10 Uhr wurde im Fundamentgraben an der nordwestlichen Ecke des Anbaus eine nicht bezeichnete Leitung mit Außendurchmesser etwa 110 mm freigelegt. Die Leitung führt diagonal durch die geplante Fundamentachse. Bei Freilegung trat kein Wasser aus. Ob sie noch in Betrieb ist, konnte unser Polier nicht feststellen.",
      ("h","1 Betroffene Arbeiten"),"Der Aushub und die Betonage im westlichen Fundamentabschnitt von Achse 2 bis 4 können bis zur Klärung nicht fortgeführt werden. Die Leitung liegt etwa 0,65 m unter Hofniveau und würde die geplante Gründung kreuzen. Wir haben den Bereich mit Brettern gesichert und nicht weiter geöffnet. Der Betonabruf für diesen Abschnitt am 07.09.2023 wurde zunächst angehalten.",
      "Betroffen sind ein Bagger mit Fahrer und zwei Facharbeiter. Zwei weitere Beschäftigte konnten in den Bestandsrückbau umgesetzt werden. Die Bauleitung wurde telefonisch um 08:35 Uhr informiert. Herr Winter erklärte vor Ort, dass er die Leitung für stillgelegt halte, konnte aber keinen Stilllegungsnachweis vorlegen. Die Suchschlitze von 2021 liegen östlich des jetzt angetroffenen Verlaufs.",
      ("h","2 Erforderliche Klärung"),"Bitte veranlassen Sie eine eindeutige Zuordnung und entscheiden Sie, ob eine Umlegung erforderlich ist. Ein Eingriff wird von uns ohne Freigabe des Eigentümers nicht vorgenommen. Die anderen zugänglichen Arbeitsbereiche werden weiterbearbeitet. Die Dauer der Unterbrechung hängt von der Klärung ab und lässt sich heute noch nicht abschließend angeben.",
      "Wir behalten uns die Abrechnung nachgewiesener zusätzlicher Leistungen und belegter Stillstandskosten vor. Diese Anzeige enthält noch kein beziffertes Nachtragsangebot. Ein Fortschreiben des Terminplans erfolgt nach Freigabe. Als Beleg verweisen wir auf den Tagesbericht vom 05.09.2023; der Leitungsfund ist dort von der Bauleitung aufgenommen. Die Folgen für den Gesamtfertigstellungstermin sind heute nicht abschließend bestimmbar."
    ],attachments=[])
    record(33,"bau","verein","Nachtragsangebot N01 Leitungsumlegung",[
      "nach der Freilegung und Abstimmung mit dem Hauswart bieten wir die Umlegung des privaten alten Pumpenstrangs an. Der Verein hat den Strang am 11.09.2023 dem eigenen Grundstücksbetrieb zugeordnet. Er führt weiterhin Wasser aus dem Nebengebäude ab und ist nicht stillgelegt. Der neue Verlauf liegt außerhalb der Fundamentachse.",
      ("h","1 Zusätzliche Bauleistungen"),("t",[["Position","Menge","EP netto EUR","Betrag EUR"],["N01.1 Rohrleitung mit Formteilen","14,00 m","145,00","2.030,00"],["N01.2 Zusätzlicher Graben","4,00 m³","92,00","368,00"],["N01.3 Abtransport Altmaterial","1 psch","380,00","380,00"],["N01.4 Anschlussarbeiten","16,00 h","58,00","928,00"],["Bauleistungen netto","","","3.706,00"],["N01.5 Gerätebereitschaft","2 Tage","620,00","1.240,00"],["Angebot gesamt netto","","","4.946,00"],["Umsatzsteuer 19 Prozent","","","939,74"],["Angebot gesamt brutto","","","5.885,74"]],[230,70,95,100]),
      ("h","2 Ausführung und Nachweise"),"Die Umlegung kann nach Auftrag am 18. und 19.09.2023 erfolgen. Die Anschlussarbeiten enthalten zwei Beschäftigte für jeweils acht Stunden. Die Leitung wird vor Verfüllung auf freien Durchgang und sichtbare Undichtigkeiten geprüft. Der neue Verlauf wird im Aufmaß verzeichnet. Das Rohrmaterial wird nach Auftrag disponiert.",
      "Die Gerätebereitschaft bezieht sich auf den 06. und 07.09.2023. Der Bagger stand auf der Baustelle; der Fahrer wurde teilweise mit Sicherungsarbeiten beschäftigt. Wir setzen hierfür jeweils 620,00 EUR an. Der Betrag wird unabhängig von den zusätzlichen Bauleistungen ausgewiesen, weil der Verein die Stillstandsdauer noch prüft. Die vollständige Rückbaukolonne war nicht durchgehend untätig.",
      ("h","3 Abgrenzung"),"Die zusätzlichen vier Kubikmeter Graben sind nicht in der späteren Fundamentaushubmenge der Position 2.3 enthalten. Der laufende Fundamentaushub wird separat aufgemessen. Eine doppelte Abrechnung erfolgt nicht. Arbeiten zur Umnutzung des Vereinsraums sind nicht Gegenstand dieses Nachtrags. Bitte bestätigen Sie Bauleistungen und Gerätebereitschaft gesondert. Wir halten die Preise bis 22.09.2023 aufrecht."
    ],attachments=[31])
    record(34,"verein","bau","Auftrag N01 ohne Gerätebereitschaft",[
      "Frau Döring und Herr Feldmann bestätigen gemeinsam den Auftrag für die Positionen N01.1 bis N01.4 Ihres Angebots vom 12.09.2023 mit 3.706,00 EUR netto zuzüglich 704,14 EUR Umsatzsteuer, insgesamt 4.410,14 EUR brutto. Bitte führen Sie die Umlegung wie angeboten am 18. und 19.09.2023 aus und dokumentieren Sie den neuen Verlauf vor Verfüllung.",
      "Die Position N01.5 über 1.240,00 EUR netto beauftragen oder anerkennen wir nicht. Frau Hartwig prüft hierzu noch die Tagesberichte und die Einsatzmöglichkeiten im Bestandsbereich. Herr Winter berichtet, dass am 06.09.2023 weiter abgebrochen wurde. Daraus ziehen wir heute noch keine abschließende Aussage über Ihren Geräteeinsatz. Bitte legen Sie die Einsatzaufzeichnungen vor.",
      "Der Auftrag umfasst keine Änderung der Vereinsraumnutzung. Zusätzliche Leitungsgrabenmengen dürfen nicht zugleich als Fundamentaushub abgerechnet werden. Die Bauleitung darf die Maße und den technischen Anschluss kontrollieren; eine darüber hinausgehende Vergütungszusage ist damit nicht verbunden. Im Anhang befindet sich das vollständige Nachtragsangebot als Bezug. Die nicht bestätigte Geräteposition bleibt sichtbar und wurde von uns nicht aus Ihrer Datei entfernt."
    ],attachments=[33])
    formulas=["1 Baustelleneinrichtung","18,00 × 12,00","60,00 + 6,00","24,00 + 1,20","10,00 × 8,00","95,00 + 3,00","36,00 + 6,00","144,00 + 6,00","10,00 × 8,00","3 Außentüren","12,00 × 1,50 + 1,50 × 1,50 + 1,20","9,00 + 0,40"]
    record(35,"plan","bau","Gemeinsames Aufmaß AM-07",[
      "am 16.02.2024 nahmen Lena Hartwig und Torsten Albrecht von 08:30 bis 11:15 Uhr die Schlussmengen des Loses 1 auf. Die bereits verdeckten Fundamente wurden anhand der Aufnahmen vor Betonage vom 20.09.2023 abgeglichen. Die nachstehende Mengenfeststellung bestätigt keine zusätzliche Vergütungsabrede.",
      ("h","1 Mengen und Rechenwege"),("t",[["Position","Rechenweg","Menge","Einheit"]]+[[f"2.{i+1} {ITEMS[i][0]}",formulas[i],str(FINAL_QTY[i]).replace('.',','),UNITS[i]] for i in range(12)],[193,187,65,50]),
      ("h","2 Abgrenzungen"),"Bei Position 2.3 sind nur die Fundamentgräben einschließlich der örtlichen Verbreiterung erfasst. Die vier Kubikmeter Leitungsgraben aus N01 werden nicht nochmals angesetzt. Das Fundamentvolumen erhöht sich wegen der Verbreiterung um 1,20 m³. Die zusätzliche Sockelfläche von 6 m² wurde vor Wiederverfüllung am 03.10.2023 gemeinsam festgestellt; eine Horizontalsperre wurde nicht ausgeführt.",
      "Die Rampenfläche enthält die zwei Läufe mit 18,00 m², das Zwischenpodest mit 2,25 m² und einen tatsächlich betonierten zusätzlichen Anschlussstreifen von 1,20 m². Die Rinne ist einschließlich des angepassten Endstücks 9,40 m lang. Die Türhöhe am Hauptzugang wurde gegenüber dem Bestandsbezug vor Ort kontrolliert. Eine Messung des Entwässerungsgefälles ist nicht Bestandteil dieser Mengentabelle.",
      ("h","3 Nachtragsmengen"),"Die Parteien bestätigen 14 m Leitung, 4 m³ gesonderten Leitungsgraben, einen Abtransport und 16 Anschlussstunden für N01.1 bis N01.4. Zur Gerätebereitschaft N01.5 werden keine Mengen gemeinsam bestätigt. Herr Albrecht hält seinen Ansatz von zwei Tagen aufrecht. Frau Hartwig verweist auf die offenen Einsatznachweise.",
      "Herr Albrecht und Frau Hartwig zeichnen die Mengentabelle. Der Bauherr erhält eine Kopie. Die Aufmaßunterzeichnung ersetzt weder die Abnahme noch die Entscheidung über die streitige Gerätebereitschaft. Die im Bautagebuch dokumentierten Arbeitsunterbrechungen bleiben unverändert Bestandteil der Bauakte."
    ],"Lena Hartwig, Objektüberwachung; Torsten Albrecht, Leinebau")
    base=round(sum(q*p for q,p in zip(FINAL_QTY,LEINE_EP)),2); claim=base+4946; checked=base+3706
    record(36,"bau","verein","Schlussrechnung LB-240301",[
      "für die vom 21.08.2023 bis 29.02.2024 ausgeführten Leistungen des Loses 1 berechnen wir die nachstehenden Beträge. Rechnungsnummer LB-240301, Projekt BH-21-04, Auftrag vom 18.04.2023. Leistungsempfänger ist der Bürgerhaus Leinewinkel e.V., Lindengasse 18, 37574 Einbeck. Die endgültige Fertigmeldung erfolgt nach Abschluss der noch ausstehenden Einstellarbeiten an den Türen.",
      ("h","1 Abgerechnete Leistungen"),("t",[["Position","Menge","EP netto EUR","Betrag EUR"]]+[[f"2.{i+1} {ITEMS[i][0]}",str(q).replace('.',',')+" "+UNITS[i],euro(LEINE_EP[i]),euro(q*LEINE_EP[i])] for i,q in enumerate(FINAL_QTY)]+[["Grundauftrag nach Aufmaß","","",euro(base)],["N01.1 bis N01.4","gemäß Auftrag","",euro(3706)],["N01.5 Gerätebereitschaft","2 Tage","620,00",euro(1240)],["Summe netto","","",euro(claim)],["Umsatzsteuer 19 Prozent","","",euro(claim*.19)],["Gesamt brutto","","",euro(claim*1.19)]],[217,92,86,100]),
      ("h","2 Bereits vereinnahmte Abschläge"),("t",[["Zahlung","Netto EUR","Umsatzsteuer EUR","Brutto EUR"],["12.10.2023 zu LB-231002","30.000,00","5.700,00","35.700,00"],["18.12.2023 zu LB-231201","40.000,00","7.600,00","47.600,00"],["Summe vereinnahmt","70.000,00","13.300,00","83.300,00"],["Noch verlangt",euro(claim-70000),euro((claim-70000)*.19),euro(claim*1.19-83300)]]),
      "Die Rechnung geht am 04.03.2024 beim Verein ein; wir bitten um Zahlung binnen 21 Kalendertagen nach Zugang. Die Gerätebereitschaft wird trotz des Vorbehalts im Auftrag beansprucht. Wir verweisen auf die Behinderungsanzeige und den Bauablauf. Eine abschließende Einigung hierzu liegt nicht vor. Die übrigen Positionen beruhen auf AM-07.",
      "Zahlungsempfänger ist Leinebau Hochbau GmbH. Die Zahlung erfolgt auf das im Vertrag hinterlegte Geschäftskonto. Steuernummer 35/200/48160. Geschäftsführer Henning Seidel, Am Werkhof 6, 37574 Einbeck. Für Rückfragen steht Torsten Albrecht zur Verfügung."
    ],"Henning Seidel, Geschäftsführer",[35,33])
    record(38,"verein","bau","Abnahmeprotokoll Los 1",[
      "am 18.03.2024 wurde das Los 1 zwischen 10:00 und 12:10 Uhr gemeinsam begangen. Teilgenommen haben Maren Döring und Uwe Feldmann für den Bauherrn, Lena Hartwig für die Objektüberwachung sowie Torsten Albrecht für Leinebau. Gegenstand ist allein der Bauvertrag vom 18.04.2023 einschließlich des bestätigten Nachtrags N01.1 bis N01.4.",
      ("h","1 Erklärung des Bauherrn"),"Der Bauherr nimmt die Leistungen des Loses 1 mit den nachstehenden ausdrücklich vorbehaltenen Mängeln ab. Die Architektin empfiehlt die Abnahme unter Dokumentation dieser Restpunkte. Die Erklärung umfasst keine Leistungen anderer Unternehmer und keine Abnahme der Architektenleistungen. Hinsichtlich der streitigen Gerätebereitschaft wird keine Vergütungsentscheidung getroffen.",
      ("h","2 Festgestellte Punkte"),("t",[["Kennung","Feststellung","Vereinbarung"],["M-01","Rinnenrost am nördlichen Zugang klappert bei Überfahrt; Befestigung fehlt am letzten Element.","Leinebau befestigt bis 28.03.2024."],["M-02","Tür T-03 schließt bei geringem Öffnungswinkel nicht vollständig selbsttätig.","Leinebau stellt Schließer bis 28.03.2024 ein."],["B-01","Nordwand des Saals wirkt nach Neuverputz trocken; alter Sockelbereich bleibt aufgrund Vorgeschichte beobachtungsbedürftig.","Hauswart dokumentiert Veränderungen. Keine Aussage über verdeckte Schichten."]],[48,259,188]),
      "Die Rechte wegen M-01 und M-02 werden vorbehalten. Am Begehungstag war es trocken; eine Beregnungsprüfung am Höhenanschluss fand nicht statt. Die Begehung bestätigt daher nicht die Funktion sämtlicher Anschlüsse bei Starkregen. Frau Hartwig hat keine sichtbare Putzablösung festgestellt. Herr Albrecht erklärt, dass keine Horizontalsperre hergestellt wurde, wie bereits im Angebot abgegrenzt.",
      ("h","3 Unterlagen und Termine"),"Leinebau übergibt die Produktnachweise zu Türen und Dach sowie den eingetragenen Leitungsverlauf. Das Aufmaß AM-07 ist bereits übergeben. Die Arbeiten waren am 15.03.2024 zur Abnahme gemeldet; die Abnahme wurde auf Wunsch des Vorstands am 18.03.2024 durchgeführt. Eine Vertragsstrafe ist nicht vereinbart.",
      "Für Los 1 wird als vertraglicher Ausgangspunkt der fünfjährigen Mängelfrist die heutige Abnahme notiert. Die Fristen der anderen Gewerke und die gesonderte Abnahme der Architektenleistungen sind hiervon unabhängig. Die Zahlungsprüfung vom 11.03.2024 bleibt bestehen. Die Unterschriften bestätigen die vorstehenden Erklärungen und die gemeinsame Aufnahme, nicht den Verzicht auf vorbehaltene Rechte."
    ],"Maren Döring und Uwe Feldmann, Bauherr; Torsten Albrecht, Leinebau; Lena Hartwig, Protokoll",[35])
    record(39,"plan","verein","Übergabe an den Gebäudebetrieb",[
      "am 05.04.2024 wurde das umgebaute Bürgerhaus dem Verein und dem Hauswart Ralf Winter zur Nutzung übergeben. Die Übergabe fand von 09:00 bis 11:30 Uhr statt. Die Nutzung ist auf die genehmigte Betriebsweise und die nachstehenden Einweisungen bezogen. Sie ist keine Sammelabnahme sämtlicher Gewerke.",
      ("h","1 Übergebene Einrichtungen"),"Herr Winter erhält sechs Hauptschlüssel, zwei Technikschlüssel und zwei Schlüssel für den Putzmittelschrank. Die Schlüsselnummern BH-01 bis BH-06 sind im Schlüsselbuch eingetragen. Hauptschalter, Heizungsabsperrung, Wasserabsperrung und die Reinigungsöffnung der Entwässerungsrinne wurden vor Ort gezeigt. Die Rinne ist nach Laubfall und bei erkennbarem Rückstau zu reinigen; Arbeiten an elektrischen Einrichtungen bleiben Fachkräften vorbehalten.",
      "Die Ausgänge und Bewegungsflächen bleiben frei. Die genehmigte Gesamtbelegung beträgt 120 Personen. Der Vereinsraum bleibt Besprechungsraum. Der im August 2023 angesprochene Gymnastikbetrieb ist weder beauftragt noch bauaufsichtlich geklärt worden und wird zur heutigen Übergabe nicht aufgenommen. Die übrigen Veranstaltungen dürfen nach der genehmigten Organisation stattfinden.",
      ("h","2 Abnahmen und Restpunkte"),("t",[["Leistung","Abnahme laut Gewerkeunterlage","Stand bei Übergabe"],["Los 1 Leinebau","18.03.2024","M-01 und M-02 am 28.03.2024 nachkontrolliert; Rost befestigt, Tür eingestellt."],["Heizung und Sanitär Mertens Installation","22.03.2024","Einweisung und Funktionsprotokoll vom 22.03.2024 übergeben."],["Elektro Lichtpfad","25.03.2024","Prüfprotokoll und Verteilerbeschriftung übergeben."],["Fenster und Innenausbau Weserwerk","27.03.2024","Keine offenen Punkte im zugehörigen Abnahmeblatt."],["Außenanlagen Hofraum","02.04.2024","Pflege der Fugen und Abläufe erläutert."]],[148,142,205]),
      ("h","3 Planungsunterlagen und Betrieb"),"Übergeben werden die Genehmigung mit Bauvorlagen, die Revisionsmappe R-01, die Bedienunterlagen und das Dokumentenregister. Die Leitungsumlegung ist darin verzeichnet. Der Höhenanschluss wurde aus den Baustellenangaben übernommen. Eine zerstörende Prüfung der verdeckten Abdichtung erfolgte bei Übergabe nicht.",
      "Die Planungs- und Überwachungsleistungen der Architektin bis einschließlich Objektüberwachung werden dem Vorstand in einem gesonderten Termin zur Abnahme vorgestellt. Heute wird lediglich der Erhalt der Unterlagen bestätigt. Die Objektbetreuung läuft weiter. Eine Begehung im September 2026 wird im Kalender vorgemerkt; sie ersetzt nicht die später erforderliche Überwachung der gewerkeweisen Fristen.",
      "Herr Winter bestätigt die Einweisung und den Unterlagenerhalt. Frau Döring übernimmt die Organisation der Nutzungsbuchführung. Die Kostenübersicht ist als Stand der Rechnungsprüfung zu lesen; der Streit über die Gerätebereitschaft ist dadurch nicht erledigt."
    ],"Lena Hartwig; Maren Döring; Ralf Winter, Hauswart",[17,38])
    record(41,"verein","plan","Feuchteflecken und Wasser am Eingang",[
      "bei der Vorbereitung des Seniorennachmittags am 07.09.2026 hat Herr Winter erneut abblätternde Farbe und einen feuchten Rand an der nördlichen Saalwand festgestellt. Betroffen ist die Ecke links vom Fenster, nicht der neue Technikraum. Das beigefügte Foto wurde von Herrn Winter gestern um 10:18 Uhr aufgenommen. Es zeigt die zugängliche Oberfläche; wir haben den Putz nicht geöffnet.",
      "Außerdem blieb nach dem Regen vom 04.09.2026 Wasser unmittelbar vor der Eingangstür stehen. Die Rinne war am 03.09.2026 gereinigt worden, allerdings gibt es dazu nur den Kalendereintrag des Hauswarts. Beim Kehren ist ihm aufgefallen, dass der Rost am letzten Element etwas tiefer liegt. Ob sich das Pflaster abgesenkt hat oder die Rinne selbst, können wir nicht sagen.",
      "Bitte nehmen Sie diese Punkte im Rahmen der vereinbarten Objektbetreuung auf und laden Sie Leinebau zur Begehung ein. Die Feuchte im Altbau war schon vor dem Umbau ein Thema; für uns ist unklar, ob sie mit dem neuen Anschluss zusammenhängt. Der Vereinsraum wird weiter nur für Besprechungen genutzt. Eine Forderungssumme können wir derzeit nicht nennen. Die Veranstaltungen laufen mit freigehaltenem Wandbereich weiter; bei erneutem Wassereintritt sperrt Herr Winter den betroffenen Zugang und informiert uns."
    ],attachments=[42])
    record(43,"plan","verein","Objektbegehung am 17.09.2026",[
      "am 17.09.2026 haben Lena Hartwig, Maren Döring, Ralf Winter und Torsten Albrecht die gemeldeten Erscheinungen von 09:00 bis 11:20 Uhr gemeinsam besichtigt. Die Einladung war am 09.09.2026 telefonisch abgestimmt worden. Es war trocken, die Außentemperatur betrug etwa 17 Grad. Die Feststellungen beschränken sich auf die zugänglichen Bauteile.",
      ("h","1 Befunde an der Nordwand"),"Im Saal zeigt der nördliche Sockelbereich links vom Fenster eine unregelmäßige Verfärbung mit abblätternder Beschichtung. Die sichtbare Ausdehnung beträgt entlang der Wand etwa 2,40 m und in der Höhe maximal 0,35 m. Das Foto vom 07.09.2026 ist diesem Ort zugeordnet. An der Wandfläche oberhalb 0,60 m waren keine Ablösungen erkennbar. Es fand keine Materialprobe und keine Öffnung der Sockelabdichtung statt.",
      "Relative kapazitive Anzeige am unteren Sockel: 86 bis 92 Geräteeinheiten, Vergleichsfläche bei 1,20 m Höhe: 34 bis 39. Die Werte sind keine Masseprozente und zwischen Geräten nicht vergleichbar. Der Befund zeigt einen örtlichen Unterschied, belegt aber nicht für sich allein eine bestimmte Wasserquelle. Herr Albrecht verweist auf die schon 2021 dokumentierte Feuchte und die nicht beauftragte Horizontalsperre.",
      ("h","2 Rampe und Rinne"),"Am letzten Rinnenelement vor T-03 wurde gegenüber dem angrenzenden Belag ein Höhenversatz von etwa 6 mm gemessen. Der Rost saß bei der Begehung fest. Ein vorsichtiger Eimertest mit 10 Litern Wasser zeigte einen langsamen Ablauf und eine Restpfütze am seitlichen Anschluss. Der unterirdische Anschluss wurde nicht geöffnet. Im Rinnenkasten lag wenig feines Sediment. Herr Winter erklärte, er habe vor dem Regen Laub entfernt, den Anschlussstutzen aber nicht gespült.",
      ("h","3 Fachliche Bewertung und offener Untersuchungsbedarf"),"Das Schadensbild der Nordwand ist mit Feuchteeinwirkung vereinbar. Als Ursachen kommen verbliebene Bestandsfeuchte, eine unzureichende Trennung am Sockelanschluss oder ein Wasserzutritt im Zusammenhang mit dem Höhenanschluss in Betracht. Der jetzige Sichtbefund erlaubt keine belastbare Zuordnung. Das frühere Feuchtebild und die ausgeführte vertikale Abdichtung müssen bei einer Bauteilöffnung gemeinsam betrachtet werden.",
      "Der Höhenversatz und die verlangsamte Entwässerung sind technisch aufklärungsbedürftig. Die Rinne gehört zu Los 1, die anschließende Grundstücksleitung zu den Außenanlagen. Für die Verantwortungsabgrenzung sind Sohlenlage, Anschlussgefälle und Durchgängigkeit zu prüfen. Eine Sanierungsvariante, ein Kostenbetrag oder eine Haftungsquote werden heute nicht festgelegt.",
      ("h","4 Erklärungen der Beteiligten"),"Leinebau erklärt seine Bereitschaft zur gemeinsamen Untersuchung, erkennt aber eine eigene Verursachung nicht an. Frau Döring verlangt eine getrennte Dokumentation der Befunde vor jeder Öffnung. Frau Hartwig hält die Beweissicherung der Höhen und des Anschlusses für erforderlich. Der Verein entscheidet über Umfang und Beauftragung nach Eingang eines Untersuchungsvorschlags. Bis dahin hält Herr Winter den Rinnenbereich sauber und meldet Veränderungen."
    ],"Lena Hartwig, Architektin; Maren Döring; Ralf Winter; Torsten Albrecht",[41,42,38])
    record(44,"plan","verein","Feststellung zur laufenden Gewährleistung",[
      "im Anschluss an die Begehung vom 17.09.2026 erhalten Sie die fortgeschriebene gewerkeweise Übersicht und die fachliche Zuordnung der offenen Punkte. Die Übersicht dokumentiert den heutigen Aktenstand. Sie entscheidet weder den Streit um Verursachung noch rechtliche Fragen zu einer Veränderung des Fristlaufs.",
      ("h","1 Gewerkeweise Ausgangsdaten"),("t",[["Gewerk und Beleg","Abnahme","Vertragsdauer / rechnerisches Ende"],["Leinebau, Vertrag 18.04.2023 und Abnahmeprotokoll","18.03.2024","5 Jahre / 18.03.2029"],["Heizung und Sanitär, Register R-01 Nr. 2","22.03.2024","5 Jahre / 22.03.2029"],["Elektro, Register R-01 Nr. 3","25.03.2024","5 Jahre / 25.03.2029"],["Fenster und Ausbau, Register R-01 Nr. 4","27.03.2024","5 Jahre / 27.03.2029"],["Außenanlagen, Register R-01 Nr. 5","02.04.2024","5 Jahre / 02.04.2029"],["Architektur, Vertrag 18.01.2021","Keine gesonderte Erklärung in der vorliegenden Mappe","Fristbeginn nicht aus der Gebäudeübergabe übernommen"]],[200,88,207]),
      "Die Daten 2029 sind ausschließlich rechnerische Kontrolltermine aus den dokumentierten Vertragsangaben, keine künftig eingetretenen Ereignisse. Die Registerangaben der anderen Gewerke sind vor einer Fristentscheidung mit den dort bezeichneten Originalen abzugleichen. Die Mängelmeldung vom 08.09.2026 wird nicht als automatische Verlängerung aller Fristen verbucht. Die Vereinbarung eines späteren Untersuchungstermins allein wird ebenfalls nicht als abschließende rechtliche Aussage zum Fristlauf behandelt.",
      ("h","2 Technischer Stand"),"Die 2024 bei Abnahme festgehaltenen Punkte M-01 und M-02 wurden laut Übergabeprotokoll am 28.03.2024 nachkontrolliert. Das jetzige Wasserbild am Eingang ist neu aufzunehmen; der feste Rost am Begehungstag beweist nicht die ausreichende Entwässerung. Der Wandbefund liegt im alten Saal und kann nicht ohne Untersuchung dem Anbau oder allein dem Bestand zugerechnet werden. Die angebotene Sockelabdichtung und die ausgeschlossene Horizontalsperre sind bei der technischen Prüfung auseinanderzuhalten.",
      ("h","3 Weitere Betreuung und Sicherheiten"),"Der Verein hat noch keinen Auftrag zur Bauteilöffnung erteilt. Leinebau hat kein Anerkenntnis abgegeben. Die Architektin hält die getrennte Aufnahme von Höhenanschluss, Rohrdurchgang und Wandaufbau für notwendig, bevor ein Sanierungsumfang bestimmt wird. Die Rückmeldung der Beteiligten ist offen. Nach dem Bauvertrag Los 1 besteht keine Gewährleistungsbürgschaft und kein pauschaler Gewährleistungseinbehalt; insoweit ist heute nichts freizugeben.",
      "Die nächste Fristenkontrolle wird für September 2027 im Büro vorgemerkt. Eine weitere Begehung rechtzeitig vor den jeweiligen gewerkeweisen Endterminen bleibt vertraglich vorgesehen. Diese Vormerkung dokumentiert nur einen künftigen Arbeitsbedarf, keinen bereits eingetretenen Schaden. Dem Vorstand wird die vollständige Begehungsniederschrift übergeben; eine abschließende Sanierungs- oder Haftungsentscheidung enthält die heutige Feststellung nicht."
    ],attachments=[43,39,40])

def line_text(c,x,y,text,size=10,bold=False):
    c.setFillColor(colors.black);c.setFont("TNRB" if bold else "TNR",size);c.drawString(x,y,text)

def wrapped(c,x,y,text,width=300,size=10):
    p=para(text,ParagraphStyle("Plantext",parent=BODY,fontSize=size,leading=size+3))
    _,height=p.wrap(width,1000);p.drawOn(c,x,y-height);return y-height-10

def dim(c,x1,y1,x2,y2,label):
    c.setStrokeColor(colors.HexColor("#60666a"));c.setLineWidth(.5);c.line(x1,y1,x2,y2)
    for x,y in [(x1,y1),(x2,y2)]:c.line(x-3,y-3,x+3,y+3)
    line_text(c,(x1+x2)/2-12,(y1+y2)/2+5,label,9)
    c.setStrokeColor(colors.black)

PLAN_NOTES = {
 4:("B-00", "Bestandsaufmaß und Suchstellen", "10.02.2021",[
   "Aufgemessen durch Lena Hartwig und Ralf Winter am 10.02.2021. Bestand außen 18,00 m mal 12,00 m, Bruttofläche 216 m². Nordpfeil und Maßketten beziehen sich auf das gesamte Bestandsgebäude. Räume: Saal 120 m², Vereinsraum 24 m², Küche 12 m², Lager 10 m² und Flur 20 m²; Wand- und Konstruktionsfläche 30 m².",
   "Suchschlitz S1 liegt mittig an der Nordwand, S2 an der nordöstlichen Ecke. Beide wurden bis zur Fundamentoberkante geöffnet. Dort war eine Außenabdichtung nur in Resten erkennbar. Die genaue Lage einer Horizontalsperre konnte nicht festgestellt werden. Die nördliche Saalwand zeigt innen Verfärbung bis etwa 25 cm Höhe.",
   "Kapazitive Vergleichsanzeigen: Messpunkt M1 am Saalsockel 84, M2 an der Nordostecke 79, M3 am südlichen Sockel 38 und M4 in 1,20 m Höhe an der Nordwand 35 Geräteeinheiten. Das Gerät liefert relative Vergleichswerte, keine Masseprozente. Die Aufnahmesituation war unbeheizt bei etwa 7 Grad Raumtemperatur.",
   "Der gestrichelte Leitungsverlauf ist ausschließlich die Erinnerung des Hauswarts aus der Skizze von 1992. Er ist nicht geortet und darf nicht als gesicherte Lage für Erdarbeiten verwendet werden. Eine flächige Untersuchung des Hofs wurde nicht durchgeführt. Die Variantenfelder zeigen die Lage des möglichen Nordanbaus mit 10,00 mal 8,00 m und des alternativen Westanbaus mit 6,00 mal 8,00 m; sie sind keine Ausführungsfreigabe."]),
 7:("E-01", "Entwurfsplan Nordanbau", "10.06.2021",[
   "Entwurf auf Grundlage der Variantenentscheidung vom 25.03.2021. Der 18,00 m mal 12,00 m große Bestand bleibt erhalten. Der Anbau liegt an der Nordseite, 4,00 m von der westlichen Bestandskante versetzt, und misst 10,00 m mal 8,00 m. Die Gesamtbruttofläche beträgt 296 m².",
   "Foyer 32 m², Sanitärzone 20 m², Lager 8 m² und Technik 8 m² ergeben 68 m² Nutzfläche im Anbau. Die Außen- und Innenwände beanspruchen weitere 12 m². Die im Plan eingetragenen Raumflächen sind mit der Objektbeschreibung abgestimmt. Tür T-02 verbindet den Saal mit dem Foyer.",
   "Höhenbezug: Saalfußboden 112,40 m, Hof am Rampenantritt 111,80 m. Die Rampe liegt nördlich des Foyers innerhalb des eigenen Grundstücks. Zwei Läufe von jeweils 6,00 m Länge überwinden mit 5 Prozent Gefälle je 0,30 m. Zwischenpodest 1,50 mal 1,50 m. Alle Höhen beruhen auf dem örtlichen Projektbezugspunkt, nicht auf einem amtlichen Nivellement.",
   "Das Mauerwerk des Anbaus erhält eine eigenständige Gründung. Der Bestandsdurchbruch wird mit der Tragwerksplanung abgestimmt. Die in Rot eingetragenen Bauteile sind neu, schwarze Bauteile bleiben bestehen. Küchen- und Vereinsraumnutzung werden nicht erweitert. Entwurfsfreigabe durch den Vorstand am 15.06.2021; dieses Blatt ist keine Ausführungsunterlage."]),
 12:("L-01", "Lageplan zum Bauantrag", "20.01.2022",[
   "Bauherr: Bürgerhaus Leinewinkel e.V., Lindengasse 18, 37574 Einbeck. Entwurfsverfasserin: Lena Hartwig, Mühlenweg 7, 37574 Einbeck. Baugrundstück: Gemarkung Einbeck, Flur 14, Flurstück 83/7, Fläche 1.512 m². Das Blatt ist der objektbezogene Lageplan der Entwurfsverfasserin nach den vom Bauherrn übergebenen Grundstücksangaben; es wird kein amtliches Vermessungssiegel geführt.",
   "Grundstück 42,00 m breit und 36,00 m tief. Der Bestand liegt 10,00 m von der Westgrenze und 8,00 m von der Straßengrenze entfernt. Seine Grundfläche beträgt 216 m². Der Anbau liegt 14,00 m von der Westgrenze entfernt und endet 8,00 m vor der Nordgrenze. Seine Grundfläche beträgt 80 m². Überbaute Fläche gesamt 296 m², Verhältnis zur Grundstücksfläche rund 0,196.",
   "Der Abstand des Bestands zur Ostgrenze beträgt 14,00 m. Das Grundstück wird von Süden über die Lindengasse erschlossen. Zwölf Kfz-Stellplätze liegen südlich und östlich des Gebäudes; Stellplatz 1 ist auf 3,50 m verbreitert. Acht Fahrradplätze liegen westlich des Rampenzugangs. Der Zugang zu beiden Ausgängen bleibt auf dem eigenen Grundstück.",
   "Nachbarflächen sind westlich Flurstück 83/6, nördlich 84/1 und östlich 83/8. Es ist keine Bebauung auf fremdem Grund beantragt. Die eingezeichnete Regenwasserleitung folgt dem geplanten Anschluss an den vorhandenen Grundstücksstrang; die endgültige Höhenlage wird vor Ausführung ermittelt. Grundlage der Beurteilung ist der Bauantrag vom 20.01.2022 mit Baubeschreibung B-01."]),
 13:("G-01", "Bauzeichnung Grundriss Schnitt Ansichten", "20.01.2022",[
   "Diese Bauzeichnung gehört zum Antrag 63-2022-041. Blatt 1 enthält den Grundriss des gesamten Erdgeschosses mit Nutzungsflächen, Türkennungen und Rettungswegen. Blatt 2 enthält den Gebäudeschnitt und die beiden für den Anbau maßgeblichen Ansichten. Der Bestand ist schwarz, der Anbau rot dargestellt. Sämtliche Maße sind in Metern angegeben.",
   "Hauptabmessungen: Bestand 18,00 mal 12,00 m, Anbau 10,00 mal 8,00 m. Fertigfußboden 112,40 m im örtlichen Projektbezug. Bestandstraufe 115,80 m, Bestandsfirst 117,70 m, Attika Anbau 115,90 m. Die Nutzräume liegen auf einer Ebene. Das Dachgeschoss bleibt unausgebaut und ist kein Aufenthaltsbereich.",
   "Die Türen T-01 und T-03 sind Außentüren mit 1,00 m lichter Breite. T-02 verbindet Saal und Foyer ebenfalls mit 1,00 m lichter Breite. Der freie Bestandsflur ist mindestens 1,20 m breit. Die Gesamtbelegung von 120 Personen und die Räume S-01, V-01 und F-01 entsprechen B-01 und der nachgereichten Ergänzung BS-02.",
   "Die im Schnitt gezeigte Bodenplatte des Anbaus liegt mit ihrer Rohoberkante 0,16 m unter Fertigfußboden. Das Flachdach fällt vom Bestand weg nach Norden. Die Bewegungsfuge wird durch die Gebäudehülle geführt. Die nördliche Eingangstür wird stufenfrei an die Rampe angeschlossen. Konstruktions- und Bewehrungsdetails sind in der nachfolgenden Ausführungsplanung festzulegen."]),
 20:("AP-01", "Ausführungsplan Erdgeschoss", "02.06.2022",[
   "Freigegebene Gebäudefassung Revision 00 vom 02.06.2022. Grundlage sind Genehmigung 63-2022-041, G-01 und Tragwerksplan T-02. Das Blatt legt Wandstärken, Rohbauöffnungen, Bewegungsfuge und Bezugshöhen für Los 1 fest. Maße vor Bestellung von Bauteilen örtlich prüfen. Abweichungen der Fachplanerblätter sind vor Ausführung über die Objektplanung zu klären.",
   "Anbau außen 10,00 mal 8,00 m. Außenwand 24 cm Mauerwerk mit 16 cm Mineralwolldämmung und Putz. Bewegungsfuge zum Bestand 30 mm; keine starre Mörtelbrücke. Rohbauöffnung T-02 1,135 mal 2,135 m, T-03 1,135 mal 2,135 m. Fertige lichte Durchgangsbreite jeweils 1,00 m. Die Türanschlüsse werden am Rohbau gemeinsam aufgemessen.",
   "Rohoberkante Bodenplatte 112,24 m, Fertigfußboden 112,40 m. Bodenaufbau 16 cm einschließlich Estrich und Belag. Fundamentbreite 50 cm, Höhe 80 cm gemäß T-02. Die Leitungsdurchführungen liegen in Hülsen und nicht im Auflagerbereich des Bestandssturzes. Die unbekannte Hofleitung ist vor Aushub zu lokalisieren; der gestrichelte Altverlauf ist kein Leitungsfreigabevermerk.",
   "Detail D-12 regelt den Höhenanschluss. Solloberkante Rinnenrost 112,38 m, angrenzender Belag am Eingang 112,40 m. Die Oberfläche fällt zur Rinne. Das Dachgefälle ist mit 2 Prozent nach Norden vorgesehen. Vor Verfüllung sind Sockelabdichtung und Schutzlage durch die Objektüberwachung zu kontrollieren. Die Portalmarkierung einer technischen Koordinationsdatei ist keine Revision dieses freigegebenen Gebäudeplans."]),
 21:("D-12", "Anschlussdetail Rampe Rinne Altbausockel", "06.06.2022",[
   "Detailblatt D-12 Revision 00, freigegeben durch Lena Hartwig am 06.06.2022. Das Detail gehört zu AP-01 und zum Leistungsverzeichnis Los 1. Es zeigt den örtlichen Eingangsanschluss im Schnitt; die Rampenlängen werden im Grundriss bestimmt. Es ist kein flächiger Sanierungsplan für den gesamten Altbau.",
   "Der Fertigfußboden liegt bei 112,40 m. Der Rinnenrost unmittelbar vor dem Zugang erhält Sollhöhe 112,38 m. Der befestigte Anschlussbereich fällt auf einem Meter Länge um 2 cm zur Rinne. Die Türschwelle bildet keinen aufstehenden Absatz. Der Rinnenkasten ist gegen den Türanschluss abzudichten und am ersten Stutzen an die Grundstücksleitung anzuschließen.",
   "Die zweilagige vertikale Sockelabdichtung wird auf tragfähigem gereinigtem Untergrund aufgebracht und mindestens bis 30 cm über anschließendes Gelände geführt, soweit keine Türöffnung liegt. Am niveaugleichen Türanschluss wird sie an den hierfür geeigneten Anschlussflansch geführt. Die Schutzlage darf beim Verfüllen nicht verrutschen. Der Übergang zur Bestandswand ist vor dem Verdecken fotografisch und im Bautagebuch festzuhalten.",
   "Die Abdichtung im Sockelbereich ersetzt keine nachträgliche Horizontalsperre im Altbau. Vor Neuverputz sind die Befunde im Innenbereich aufzunehmen. Zuständig für Rinne und direkten Anschluss ist Los 1; ab dem ersten Rohrstutzen liegt die Leistung bei den Außenanlagen. Die Durchgängigkeit ist gewerkeübergreifend vor Übergabe zu prüfen. Geänderte Höhen aus Haustechnikblättern bedürfen einer ausdrücklichen Revision dieses Details."]),
}

def draw_building(c, number):
    u=72/25.4*1000/150; x,y=(158 if number==4 else 68),107
    def rect(mx,my,mw,mh,fill=None):
        c.setFillColor(colors.HexColor(fill) if fill else colors.white)
        c.rect(x+mx*u,y+my*u,mw*u,mh*u,fill=bool(fill),stroke=1)
    c.setLineWidth(2);rect(0,0,18,12)
    c.setLineWidth(.8);c.line(x+12*u,y,x+12*u,y+12*u);c.line(x+12*u,y+6*u,x+18*u,y+6*u)
    c.line(x+12*u,y+9*u,x+18*u,y+9*u);c.line(x+15*u,y+6*u,x+15*u,y+12*u)
    c.setFillColor(colors.black)
    for xx,yy,label in [(1,5,"S-01 Saal 120 m²"),(12.4,2.8,"V-01 24 m²"),(12.3,7.4,"Küche"),(15.2,7.4,"Lager"),(12.4,10.4,"Flur 20 m²")]:
        line_text(c,x+xx*u,y+yy*u,label,9)
    if number !=4:
        c.setStrokeColor(colors.HexColor("#9c3434"));c.setLineWidth(2);rect(4,12,10,8)
        c.setLineWidth(.8);c.line(x+10*u,y+12*u,x+10*u,y+20*u)
        c.line(x+10*u,y+16*u,x+14*u,y+16*u);c.line(x+12*u,y+12*u,x+12*u,y+16*u)
        c.setFillColor(colors.black)
        for xx,yy,label in [(4.5,16,"F-01 32 m²"),(10.3,18,"WC 20 m²"),(10.2,14,"L 8"),(12.2,14,"T 8")]:
            line_text(c,x+xx*u,y+yy*u,label,9)
        c.setStrokeColor(colors.black)
        for xx,yy,lab in [(6,0,"T-01"),(7,12,"T-02"),(6,20,"T-03")]:
            c.setStrokeColor(colors.white);c.setLineWidth(5);c.line(x+xx*u,y+yy*u,x+(xx+1)*u,y+yy*u)
            c.setStrokeColor(colors.black);c.setLineWidth(.7);c.line(x+xx*u,y+yy*u,x+xx*u,y+(yy+1)*u)
            line_text(c,x+(xx+1.3)*u,y+yy*u+4,lab,8)
        dim(c,x+4*u,y+20*u+20,x+14*u,y+20*u+20,"10,00")
        dim(c,x+14*u+22,y+12*u,x+14*u+22,y+20*u,"8,00")
        c.setStrokeColor(colors.HexColor("#247343"));c.setLineWidth(1.5)
        c.line(x+5*u,y+7*u,x+5*u,y+1*u);c.line(x+8*u,y+8*u,x+8*u,y+18*u)
        line_text(c,510,452,"Grün: Wege zu den Ausgängen",10)
    else:
        c.setDash(4,3);c.setLineWidth(.8);rect(4,12,10,8);c.setDash()
        line_text(c,x+4.5*u,y+15*u,"Variante 1: Nordanbau 80 m²",9)
        c.setStrokeColor(colors.HexColor("#935d22"));c.setDash(4,3)
        c.rect(x-6*u,y+2*u,6*u,8*u);c.setDash();c.setStrokeColor(colors.black)
        line_text(c,510,450,"Variante 2: Westanbau 48 m²",10)
        for mx,my,label in [(9,12,"S1"),(17,12,"S2"),(6,11,"M1"),(17,11,"M2"),(6,1,"M3"),(6,10,"M4")]:
            c.circle(x+mx*u,y+my*u,3,stroke=1,fill=0);line_text(c,x+mx*u+5,y+my*u+3,label,8)
    c.setStrokeColor(colors.black);c.setLineWidth(.6)
    dim(c,x,y-23,x+18*u,y-23,"18,00")
    dim(c,x+18*u+25,y,x+18*u+25,y+12*u,"12,00")
    c.setDash(4,3);c.line(x+7*u,y+12*u,x+15*u,y+20*u);c.setDash()
    if number==4:
        line_text(c,565,420,"Altleitung: nur Erinnerung des Hauswarts",10)
        line_text(c,565,390,"Saalfußboden 112,40 m",11)
        line_text(c,565,370,"Hof etwa 0,60 m tiefer",11)
        wrapped(c,565,335,"S1 und S2: punktuelle Suchschlitze. M1 bis M4: relative kapazitive Vergleichsanzeigen. Keine flächige Leitungssuche.",230)
        wrapped(c,565,248,"Die gestrichelten Baukörper sind untersuchte Varianten, keine Genehmigungs- oder Ausführungsplanung. Der Saalsockel ist örtlich verfärbt.",230)
        line_text(c,565,150,"Maßstab 1:150 bei A4",10)
        c.rect(565,117,5*u,6,fill=0);line_text(c,565,101,"0",9);line_text(c,565+5*u-10,101,"5 m",9)
        c.line(748,472,748,509);line_text(c,742,520,"N",12,True)
        return
    line_text(c,510,420,"Gestrichelt: ungeorteter Altleitungsverlauf",10)
    line_text(c,510,390,"Fertigfußboden 112,40 m",11)
    line_text(c,510,370,"Hof am Rampenantritt 111,80 m",11)
    line_text(c,510,350,"Anbau: Wände 12 m², Nutzung 68 m²",10)
    wrapped(c,510,320,"Rampe nördlich des Foyers: zwei Läufe je 6,00 × 1,50 m, Zwischenpodest 1,50 × 1,50 m. Höhenunterschied 0,60 m; 5 Prozent Längsgefälle.",255)
    wrapped(c,510,238,"Wandstärken und Öffnungsrohbaumaße sind in AP-01 angegeben. Rinnen- und Sockelanschluss siehe D-12. Der Grundriss stellt keine Leitungserkundung dar.",255)
    c.line(748,472,748,509);line_text(c,742,520,"N",12,True)
    line_text(c,510,145,"Maßstab 1:150 bei Druck auf A4",10)
    c.rect(510,117,5*u,6,fill=0);line_text(c,510,101,"0",9);line_text(c,510+5*u-10,101,"5 m",9)

def draw_site(c):
    u=72/25.4*1000/250;x,y=45,83
    c.setLineWidth(1);c.rect(x,y,42*u,36*u)
    c.setLineWidth(2);c.rect(x+10*u,y+8*u,18*u,12*u)
    c.setStrokeColor(colors.HexColor("#9c3434"));c.rect(x+14*u,y+20*u,10*u,8*u)
    c.setStrokeColor(colors.black);c.setLineWidth(.6)
    c.rect(x+14*u,y+28*u,8*u,3.5*u)
    line_text(c,x+14*u+3,y+29*u,"Rampe",9)
    line_text(c,x+12*u,y+13*u,"Bestand 216 m²",10)
    line_text(c,x+14.5*u,y+24*u,"Anbau 80 m²",10)
    for i in range(8):
        w=3.5 if i==0 else 2.5; sx=x+(2+(0 if i==0 else 3.5+(i-1)*2.5))*u
        c.rect(sx,y+1*u,w*u,5*u);line_text(c,sx+4,y+3*u,str(i+1),8)
    for i in range(4):
        c.rect(x+34*u,y+(8+i*2.5)*u,5*u,2.5*u);line_text(c,x+36*u,y+(9+i*2.5)*u,str(i+9),8)
    line_text(c,70,62,"Lindengasse / öffentliche Erschließung",10)
    line_text(c,65,500,"Nordgrenze / Flurstück 84/1",10)
    dim(c,x,y-15,x+42*u,y-15,"42,00")
    dim(c,x+42*u+12,y,x+42*u+12,y+36*u,"36,00")
    dim(c,x+14*u,y+28*u,x+14*u,y+36*u,"8,00")
    line_text(c,565,468,"Flur 14 · Flurstück 83/7",12,True)
    yy=435
    for txt in ["Grundstück: 1.512 m²", "Maßstab 1:250 bei A4", "Westnachbar: 83/6", "Ostnachbar: 83/8", "Zwölf Kfz-Stellplätze", "Acht Fahrradplätze westlich", "Regenwasseranschluss: Nordhof", "Keine Inanspruchnahme fremden Grunds"]:
        yy=wrapped(c,565,yy,txt,230)
    c.line(749,480,749,515);line_text(c,743,526,"N",12,True)

def draw_section(c):
    u=72/25.4*1000/150;x,y=65,292
    c.setLineWidth(1.3);c.rect(x,y,12*u,3.4*u)
    p=c.beginPath();p.moveTo(x,y+3.4*u);p.lineTo(x+6*u,y+5.3*u);p.lineTo(x+12*u,y+3.4*u);c.drawPath(p)
    c.setStrokeColor(colors.HexColor("#9c3434"));c.rect(x+12*u,y,8*u,3.5*u)
    c.setStrokeColor(colors.black);c.line(x-15,y-.6*u,x+20*u+20,y-.6*u)
    line_text(c,x,y+122,"1 Schnitt durch Bestand und Nordanbau · 1:150",12,True)
    line_text(c,x+12*u+6,y+15,"Anbau",10)
    for yy,l in [(466,"First 117,70"),(446,"Attika 115,90"),(426,"FFB 112,40"),(406,"Rohplatte 112,24"),(386,"Hof 111,80")]:
        line_text(c,510,yy,l,10)
    c.rect(65,117,10*u,3.5*u);c.rect(115,117,1.135*u,2.135*u)
    c.rect(188,139,2*u,1.2*u);line_text(c,65,201,"2 Nordansicht Anbau · 1:150",12,True)
    c.rect(378,117,8*u,3.5*u);c.rect(410,135,2*u,1.4*u)
    line_text(c,378,201,"3 Ostansicht Anbau · 1:150",12,True)
    wrapped(c,595,200,"Traufe Bestand 115,80 m. Attika Anbau 115,90 m. Der Bezugspunkt des Projekts ist örtlich festgelegt. Öffnungsmaße siehe Grundriss und Baubeschreibung.",180)

def draw_detail(c):
    u=72/25.4*1000/10; x,y=70,310
    c.setFillColor(colors.HexColor("#eeeeee"));c.rect(x,y-.16*u,250,.16*u,fill=1)
    c.setFillColor(colors.HexColor("#c8cdd1"));c.rect(x,y-.36*u,250,.20*u,fill=1)
    c.setFillColor(colors.HexColor("#fafafa"));c.rect(x,y-.48*u,250,.12*u,fill=1)
    c.setFillColor(colors.white);c.rect(320,y,22,155,fill=1)
    c.rect(360,y-.17*u,.15*u,.15*u,fill=1)
    c.setLineWidth(2);c.line(360,y-.02*u,360+.15*u,y-.02*u)
    c.line(360+.15*u,y-.02*u,360+1.15*u,y)
    c.setLineWidth(3);c.setStrokeColor(colors.HexColor("#235b81"))
    c.line(320,y-.4*u,320,y);c.line(320,y,358,y)
    c.setStrokeColor(colors.black);c.setLineWidth(.7)
    line_text(c,70,500,"Schnitt durch Tür T-03 und Rinne · Detail 1:10",13,True)
    for xx,yy,tx,ty,label in [(200,310,70,412,"Fertigfußboden 112,40 m"),(380,304,450,420,"Rost 112,38 m"),(650,309,500,368,"Belag: 2 cm auf 1,00 m zur Rinne"),(320,228,430,225,"Anschlussflansch / vertikale Abdichtung"),(190,192,70,156,"12 cm Dämmung unter 20 cm Platte"),(382,263,485,269,"Erster Rohrstutzen: Losgrenze")]:
        c.line(xx,yy,tx,ty-4);line_text(c,tx,ty,label,10)
    line_text(c,77,282,"16 cm Bodenaufbau",10);line_text(c,77,236,"20 cm Stahlbeton",10)
    wrapped(c,70,125,"Blau: Abdichtungsführung. Höhen beziehen sich auf den örtlichen Projektbezug. Rinne 15 cm lichte Breite; Anschluss an den Türflansch vor Verdecken kontrollieren. Keine Horizontalsperre im Altbau enthalten.",700)

def plan_pdf(n):
    key,title,date,notes=PLAN_NOTES[n]
    record(n,"plan","verein" if n not in (12,13) else "amt",title,notes)
    c=canvas.Canvas(str(CASE/FILES[n]),pagesize=landscape(A4));c.setAuthor(AUTHOR);c.setTitle(title)
    def head(page):
        line_text(c,36,565,"Hartwig Architektur · Bürgerhaus Leinewinkel · Einbeck",13,True)
        line_text(c,36,547,f"{key} · {title} · {date}",11)
        line_text(c,36,28,f"BH-21-04 · gez. Lena Hartwig · Mühlenweg 7, 37574 Einbeck · Blatt {page}",9)
    head(1)
    if n==12:draw_site(c)
    elif n==21:draw_detail(c)
    else:draw_building(c,n)
    c.showPage()
    p=2
    if n==13:
        head(p);draw_section(c);c.showPage();p+=1
    head(p)
    line_text(c,45,510,"1 Erläuterungen und Zeichnungsangaben",14,True)
    yy=480
    for text in notes:yy=wrapped(c,45,yy,text,740,11)
    yy=wrapped(c,45,yy,"Empfänger: "+ACTORS[DATA[n]["recipient"]][0]+". Bezug: Bauvorhaben Lindengasse 18. Anlagen: keine. Zeichnung und Erläuterungsblatt bilden einen gemeinsamen Planstand.",740,10)
    c.save()

def screenshots():
    for n in [22,32]:
        im=Image.new("RGB",(1600,1100),"#f3f5f6");d=ImageDraw.Draw(im)
        d.rectangle((0,0,1600,90),fill="#29363d")
        d.text((36,25),"BAUPLAN   /   BH-21-04   Bürgerhaus Einbeck",font=screen_font(32,True),fill="white")
        if n==22:
            d.text((40,122),"Planstandportal",font=screen_font(38,True),fill="#17252c")
            d.text((40,180),"Projektordner > Gebäude > freigegebene Pläne",font=screen_font(23),fill="#4d5c66")
            d.text((40,225),"Aufnahme: 07.03.2023, 16:42    Angemeldet: Lena Hartwig",font=screen_font(22),fill="#4d5c66")
            rows=[("Datei / Planstand","Revision","Status","Datum"),("AP-01 Gebäude Erdgeschoss","00","Freigegeben","02.06.2022"),
                  ("D-12 Rampenanschluss","00","Freigegeben","06.06.2022"),("TGA-04 Leitungen Hof","04","Zur Koordination","07.03.2023"),("G-01 Bauzeichnung","01","Genehmigungsstand","20.01.2022")]
            for i,row in enumerate(rows):
                y=290+i*76;d.rectangle((40,y,1560,y+74),fill=("#dce4e8" if i==0 else "white"))
                for x,s in zip([57,690,885,1325],row):d.text((x,y+22),s,font=screen_font(23,i==0),fill="#22303a")
            d.rectangle((40,710,1560,1010),fill="white",outline="#bcc9d1",width=2)
            d.text((62,736),"Details zu TGA-04 / Revision 04",font=screen_font(27,True),fill="#22303a")
            for i,s in enumerate(["Hochgeladen von Paul Mertens am 07.03.2023 um 15:58.","Kommentar: Leitungsstutzen geändert. Höhenbezug Eingang noch abstimmen.",
              "Planvermerk: Rinne 112,40 m. D-12 nennt Rost 112,38 m.","Freigabe Objektplanung: ausstehend. Keine Ablösung von D-12.","Verteiler: Hartwig, Vorstand. Leinebau noch nicht in diesem Ordner."]):
                d.text((62,787+i*39),s,font=screen_font(23),fill="#22303a")
        else:
            d.text((40,122),"Baustellengruppe Bürgerhaus",font=screen_font(37,True),fill="#17252c")
            d.text((40,177),"6. September 2023   /   5 Teilnehmende   /   Export 18:10",font=screen_font(23),fill="#4d5c66")
            messages=[("08:12 · Torsten Albrecht","Leitung ist noch nicht zugeordnet. Im Westgraben bleibt der Bagger stehen.","Die beiden Leute gehen heute erst einmal in den Rückbau."),
              ("08:31 · Ralf Winter","Die Leitung müsste von der alten Pumpe kommen. Ich suche den Schlüssel","zum Nebengebäude. Bitte noch nichts abtrennen."),
              ("09:05 · Lena Hartwig","Bitte Stunden und Gerätestand getrennt dokumentieren. Der Vorstand muss","die Umlegung freigeben. Ich habe dafür keine Auftragsvollmacht."),
              ("11:42 · Torsten Albrecht","Im TGA-Blatt 04 steht die Rinne bei 112,40. Ich habe D-12 noch mit 112,38.","Welcher Plan geht an den Polier?"),
              ("12:08 · Lena Hartwig","D-12 bleibt freigegeben. TGA-04 ist nur zur Koordination eingestellt.","Die Anschlusslage prüfen wir zusammen vor dem Einbau.")]
            for i,(who,a,b) in enumerate(messages):
                y=230+i*160;x=70 if i%2==0 else 210
                d.rounded_rectangle((x,y,x+1270,y+140),radius=8,fill=("white" if i%2==0 else "#dcece4"))
                d.text((x+20,y+13),who,font=screen_font(23,True),fill="#2e4d59")
                d.text((x+20,y+54),a,font=screen_font(23),fill="#182b34");d.text((x+20,y+91),b,font=screen_font(23),fill="#182b34")
        meta=PngImagePlugin.PngInfo();meta.add_text("Author",AUTHOR);meta.add_text("Title",FILES[n]);im.save(CASE/FILES[n],pnginfo=meta)

def diary_and_archive():
    rows=[
      ("2023-08-21","trocken 22 Grad","5","Baustelle eingerichtet; Haus geräumt. Zaun und Staubschutz kontrolliert. Keine Veranstaltung im Haus.","Hartwig / Albrecht"),
      ("2023-08-22","trocken 23 Grad","5","Rückbau Boden und Einbauten begonnen. Saalsockel unverändert fleckig. Keine Putzfreigabe erteilt.","Hartwig"),
      ("2023-08-24","bewölkt 21 Grad","4","Suchbereiche an Nordwand markiert. Flächige Ortung noch nicht dokumentiert. Hauswart will Altunterlagen beschaffen.","Hartwig / Winter"),
      ("2023-08-28","leichter Regen 18 Grad","5","Aushub Ostfundament begonnen. Graben gegen Oberflächenwasser gesichert. Keine Leitung an den beiden Suchstellen sichtbar.","Albrecht"),
      ("2023-09-05","trocken 20 Grad","5","Leitung im Westgraben 08:10 freigelegt. Aushub dort angehalten. Zwei Beschäftigte in Rückbau umgesetzt; drei im Grabenbereich mit Sicherung und Klärung.","Hartwig / Albrecht"),
      ("2023-09-06","trocken 23 Grad","5","Bagger am Westgraben nicht eingesetzt. Fahrer bei Sicherung und Materialtransport. Zwei Facharbeiter im Saalrückbau. Einzelne Stillstandsstunden nicht getrennt erfasst.","Albrecht"),
      ("2023-09-07","trocken 24 Grad","4","Beton für Westabschnitt abbestellt. Innenrückbau fortgesetzt. Gerät auf Baustelle. Stundenliste zur Bereitschaft noch angefordert.","Hartwig"),
      ("2023-09-11","trocken 22 Grad","3","Leitung durch Wasserprobe dem Nebengebäude zugeordnet. Hauswart bestätigt privaten Pumpenstrang. Noch keine Umlegung begonnen.","Hartwig / Winter"),
      ("2023-09-15","bewölkt 19 Grad","4","Auftrag N01.1 bis N01.4 vom Vorstand eingegangen. Geräteposition ausdrücklich nicht bestätigt. Materialabruf für Montag.","Hartwig / Albrecht"),
      ("2023-09-18","Regen 17 Grad","4","Zusätzlichen Leitungsgraben geöffnet. Menge 4 m³ separat aufgenommen. Keine Vermischung mit Fundamentaushub.","Albrecht"),
      ("2023-09-19","trocken 18 Grad","4","14 m Rohrleitung umgelegt; Anschlussarbeiten zwei Personen je acht Stunden. Sicht- und Durchlaufprobe ohne Auffälligkeit.","Hartwig / Albrecht"),
      ("2023-09-20","trocken 18 Grad","5","Gründungsabschnitt vor Betonage kontrolliert. Fundamentvolumen 25,2 m³. Standsicherheitsunterlagen liegen laut Vorlagebestätigung vom 17.08.2023 bei Bauaufsicht.","Hartwig / Steinbach"),
      ("2023-10-03","trocken 15 Grad","3","Sockelabdichtung im freigelegten Bereich 42 m². Schutzlage abschnittsweise sichtbar. Anschluss an Türbereich noch offen. Innenputz noch nicht begonnen.","Hartwig / Albrecht"),
      ("2023-10-12","bewölkt 14 Grad","5","Mauerwerk Anbau begonnen. Erste Abschlagszahlung 35.700 EUR brutto laut Vereinsbestätigung eingegangen.","Albrecht"),
      ("2023-11-02","Regen 11 Grad","4","Dachtragwerk und Anschlüsse zum Bestand kontrolliert. Bewegungsfuge vorhanden. Dachabdichtung wegen Regen verschoben.","Hartwig"),
      ("2023-11-08","trocken 10 Grad","4","Dachabdichtung abgeschlossen. Durchdringungen sichtbar geprüft. Regenwasseranschluss noch provisorisch.","Hartwig / Albrecht"),
      ("2023-11-15","trocken 8 Grad","3","Außentüren aufgemessen. Rinnenhöhe nach D-12 besprochen. TGA-04 nicht als Gebäudeplan freigegeben.","Hartwig / Albrecht"),
      ("2023-12-04","Frost 0 Grad","3","Außenarbeiten ruhen; Installationen innen fortgesetzt. Gebäudebeheizung provisorisch. Wandfeuchte nicht als Masseprozent bestimmt.","Albrecht"),
      ("2023-12-18","bewölkt 5 Grad","3","Zweite Abschlagszahlung 47.600 EUR brutto bestätigt. Rinnenmaterial eingelagert. Innenputzbereich trocken erscheinend.","Albrecht"),
      ("2024-01-16","trocken 4 Grad","4","Innenputz im Saal erneuert. Relative Sockelanzeige 57 bis 64, obere Vergleichsstelle 33 bis 37. Keine Laborprobe beauftragt.","Hartwig / Albrecht"),
      ("2024-02-06","bewölkt 8 Grad","4","Rampe und Anschlussstreifen hergestellt. Höhen vor Einbau Rinne besprochen; gesondertes Nivellementblatt nicht in Tagesmappe.","Albrecht"),
      ("2024-02-16","trocken 10 Grad","3","Schlussaufmaß AM-07 gemeinsam aufgenommen. Rinne 9,4 m; Rampenfläche 21,45 m² einschließlich Anschlussstreifen. Keine Anerkennung von N01.5.","Hartwig / Albrecht"),
      ("2024-02-29","trocken 9 Grad","2","Rinnenanschluss und Türen eingebaut. Durchlauf mit Wasser ohne sichtbaren Rückstau bei diesem Versuch. Menge und Dauer nicht notiert.","Albrecht"),
      ("2024-03-11","bewölkt 11 Grad","2","Rechnungsprüfung abgeschlossen. Geräteposition 1.240 EUR netto weiterhin ungeklärt. Einstellarbeiten an T-03 offen.","Hartwig"),
      ("2024-03-15","trocken 13 Grad","2","Fertigmeldung Los 1. Gemeinsame Abnahme für 18.03. vereinbart. Keine Behinderung mehr angezeigt.","Albrecht"),
      ("2024-03-18","trocken 12 Grad","2","Förmliche Abnahme Los 1 mit M-01 und M-02. Kein Beregnungsversuch. Saalwand optisch ohne neue Ablösung.","Hartwig / Döring / Albrecht"),
    ]
    with (CASE/FILES[30]).open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.writer(f);w.writerow(["Datum","Witterung","Beschäftigte Los 1","Tagesbericht BH-21-04","Aufgenommen und gezeichnet"]);w.writerows(rows)
    text="""Hartwig Architektur
Mühlenweg 7, 37574 Einbeck
Dokumentenregister R-01 / Bürgerhaus Leinewinkel, Lindengasse 18
Stand 08.04.2024 / Empfänger: Vorstand und Hauswart
Bezug: Übergabe vom 05.04.2024

Sehr geehrte Damen und Herren,

1 Aufbewahrung und Vollständigkeit

Die nachstehende Liste bezeichnet die beim Verein übergebenen Originalmappen.
Die Papiermappen stehen im verschlossenen Büroschrank des Vereinsraums. Eine
digitale Kopie liegt im Vereinsarchiv. Die Liste selbst ersetzt die bezeichneten
Gewerkeunterlagen nicht. Der Vorstand hat den Empfang der Mappen am 05.04.2024
bestätigt. Die unterschriebenen Abnahmeblätter der weiteren Gewerke sind in
dieser Arbeitszusammenstellung nicht nochmals enthalten.

2 Gewerkeunterlagen und Abrechnungsregister

Nr. 1: Leinebau, Los 1. Vertrag 18.04.2023, Abnahme 18.03.2024, fünfjährige
Mängelfrist ab Abnahme. Schlussrechnung LB-240301 vom 01.03.2024 und Prüfung
vom 11.03.2024. Gerätebereitschaft bleibt streitig. Kein pauschaler Einbehalt.
Nr. 2: Mertens Installation, Heizung und Sanitär. Vertrag MI-230512,
Abnahme 22.03.2024, fünf Jahre vereinbart. Schlussrechnung MI-240322:
54.200,00 EUR netto, 10.298,00 EUR Umsatzsteuer, 64.498,00 EUR brutto.
Nr. 3: Lichtpfad Elektrotechnik, Vertrag LP-230519. Abnahme 25.03.2024,
fünf Jahre vereinbart. Schlussrechnung LP-240325: 27.400,00 EUR netto,
5.206,00 EUR Umsatzsteuer, 32.606,00 EUR brutto.
Nr. 4: Weserwerk Fenster und Ausbau, Vertrag WW-230531. Abnahme 27.03.2024,
fünf Jahre vereinbart. Schlussrechnung WW-240327: 61.800,00 EUR netto,
11.742,00 EUR Umsatzsteuer, 73.542,00 EUR brutto.
Nr. 5: Hofraum Außenanlagen, Vertrag HR-230602. Abnahme 02.04.2024,
fünf Jahre vereinbart. Schlussrechnung HR-240402: 23.600,00 EUR netto,
4.484,00 EUR Umsatzsteuer, 28.084,00 EUR brutto.
Nr. 6: Bewegliche Ausstattung Raumgut, Lieferung 02.04.2024,
Rechnung RG-240402: 12.900,00 EUR netto, 2.451,00 EUR Umsatzsteuer,
15.351,00 EUR brutto. Keine Zuordnung zur Bauwerksfrist vorgenommen.
Nr. 7: Hartwig Architektur, Vertrag 18.01.2021. Gesamt vereinbart
54.000,00 EUR netto, 10.260,00 EUR Umsatzsteuer, 64.260,00 EUR brutto.
Davon sind zum Stand der Rechnungsprüfung 960,00 EUR netto für die noch
laufende Objektbetreuung nicht als erbrachte Leistung festgestellt. Eine
gesonderte Abnahmeerklärung der Planungsleistungen liegt in dieser Mappe nicht.
Nr. 8: Untersuchung und Vermessung, Rechnungen VM-210210 und VM-230817,
zusammen 4.650,00 EUR netto, 883,50 EUR Umsatzsteuer, 5.533,50 EUR brutto.
Nr. 9: Genehmigungsgebühr 63-2022-041 und Entwässerungsgebühr,
zusammen 1.780,00 EUR ohne Umsatzsteuer. Keine Bauleistung.

3 Pläne und Nachweise

Genehmigung 63-2022-041 mit L-01, G-01, B-01 und BS-02; AP-01 und D-12;
Tragwerksmappe T-02 mit Vorlagebestätigung vom 17.08.2023; Revisionsplan
der umgelegten Pumpenleitung; Produktunterlagen der Türen und Dachbahnen;
Elektroprüfung, Heizungsprobelauf und Einweisungsprotokolle.
Die Koordinationsfassung TGA-04 ist als solche archiviert. Sie ist kein
Ersatz für den freigegebenen Rampenanschluss D-12. Ein gesondertes
unterschriebenes Nivellementblatt des Rinneneinbaus befindet sich nicht in R-01.

4 Laufende Vorgänge

Die Meldung des Vereins zur gewünschten Gymnastiknutzung wurde nicht in einen
Ausführungsauftrag überführt. Der Hauswart nutzt den Raum weiterhin als
Besprechungsraum. N01.5 ist weder anerkannt noch abschließend zurückgewiesen.
Eine gemeinsame Untersuchung später gemeldeter Mängel wird über den Vorstand
koordiniert. Sämtliche Originalmappen verbleiben beim Verein.

Mit freundlichen Grüßen
gez. Lena Hartwig, Architektin
Anlagen: keine
"""
    (CASE/FILES[40]).write_text(text,encoding="utf-8")

def model_data():
    return dict(files=FILES,phases=PHASES,qty=QTY,plan_ep=PLAN_EP,leine_ep=LEINE_EP,hage_ep=HAGE_EP,
                final_qty=FINAL_QTY,units=UNITS,items=ITEMS,
                case=str(CASE),assets=str(ASSETS),author=AUTHOR,
                documents=[dict(number=n,file=f,date=f[3:13]) for n,f in FILES.items()])

def main():
    p=argparse.ArgumentParser();p.add_argument("--stage",choices=["originals","emails","model"],default="originals");args=p.parse_args()
    CASE.mkdir(parents=True,exist_ok=True);ASSETS.mkdir(parents=True,exist_ok=True)
    register_fonts();define_documents()
    if args.stage=="model":
        (ASSETS/"einbeck-model.json").write_text(json.dumps(model_data(),ensure_ascii=False,indent=2));return
    if args.stage=="emails":
        for n,d in DATA.items():
            if FILES[n].endswith(".eml"):eml(d)
        return
    for n,d in DATA.items():
        if FILES[n].endswith(".pdf"):letter_pdf(d)
        elif FILES[n].endswith(".docx"):letter_docx(d)
    for n in PLAN_NOTES:plan_pdf(n)
    screenshots();diary_and_archive()
    (ASSETS/"einbeck-model.json").write_text(json.dumps(model_data(),ensure_ascii=False,indent=2))
    print(f"Originale und Rechenmodell erzeugt: {CASE}")

if __name__=="__main__":main()
