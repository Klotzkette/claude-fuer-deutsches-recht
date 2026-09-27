#!/usr/bin/env python3
"""Belege und kanonische Zahlungsdaten des Wohnhauses Hildesheim. Autor: Klotzkette."""
from __future__ import annotations
import json
import io
import os
import re
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, timedelta
from pathlib import Path
from bauwirtschaft_hildesheim_common import ROOT, CASE, ASSETS, ACTORS, REF, euro, build_records

QA = Path('/tmp/hildesheim/finanz')
DATA = ROOT/'quality/hildesheim-achtfamilienhaus/finanzdaten.json'
RECORDS = []
INVOICES = []


def rec(n, name, dt, actor, title, sections, recipient='bauherr', attachments=()):
    d = dict(number=n, filename=f'{n:03d}_{name}', date=dt, issuer=actor,
             recipient=recipient, title=title, sections=sections, attachments=list(attachments))
    RECORDS.append(d)
    return d


def inv(n, slug, dt, actor, ident, label, kg, items, period, *, rate=.19,
        kind='Rechnung', prior=None, note='', payment=None):
    """Items are additive current-invoice increments; finals show cumulative less previous invoices."""
    net = round(sum(q * p for _, q, _, p in items), 2)
    vat = round(net * rate, 2)
    gross = round(net + vat, 2)
    filename = f'{n:03d}_{slug}.pdf'
    obj = dict(number=n, filename=filename, date=dt, issuer=actor, id=ident,
               label=label, kg=kg, net=net, vat=vat, gross=gross, rate=rate,
               period=period, kind=kind, items=items, prior=prior, payment=payment)
    INVOICES.append(obj)
    taxid = '30/145/' + str(12000 + list(ACTORS).index(actor) * 113)
    rows = [['Position / Leistung', 'Menge', 'Einheit', 'Einzel EUR', 'Netto EUR']]
    for label_, q, unit, price in items:
        rows.append([label_, euro(q), unit, euro(price), euro(q * price)])
    sections = [f'{kind} Nr. {ident}. Rechnungsempfänger ist die Steinbogen Wohnen GmbH, Wohnhof Am Steinbogen 18, 31134 Hildesheim. Bauvorhaben: Neubau mit acht Mietwohnungen am selben Ort. Projektkennung {REF}.',
                f'Leistungszeitraum: {period}. Steuernummer des leistenden Unternehmens: {taxid}.',
                ('h', '1. Abgerechnete Leistungen'), ('t', rows, [213, 49, 48, 86, 99])]
    if prior:
        pnet = round(sum(x['net'] for x in INVOICES if x['id'] in prior), 2)
        pvat = round(pnet * rate, 2)
        sections += [f'Die folgenden Beträge bilden die kumulierte Schlussabrechnung. Die vorstehende Positionsliste enthält nur die noch nicht in den Abschlagsrechnungen abgerechneten Mengen. Die bereits abgerechneten Leistungen aus {", ".join(prior)} bleiben Bestandteil der Gesamtleistung.',
                     ('t', [['Abrechnung', 'Netto EUR', 'USt EUR', 'Brutto EUR'],
                            ['Gesamtleistung einschließlich Abschlägen', euro(pnet+net), euro(pvat+vat), euro(pnet+pvat+gross)],
                            ['Bereits berechnete und vereinnahmte Abschläge', euro(pnet), euro(pvat), euro(pnet+pvat)],
                            ['Verbleibende Forderung dieser Rechnung', euro(net), euro(vat), euro(gross)]], [215, 90, 90, 100])]
    else:
        sections += [('t', [['Abrechnung', 'Betrag EUR'], ['Entgelt netto', euro(net)],
                           [('Umsatzsteuer: steuerfreie Leistung' if actor=='bank' else f'Umsatzsteuer {euro(rate*100)} %'), euro(vat)], ['Rechnungsbetrag brutto', euro(gross)]], [360,135])]
    if rate == 0:
        sections.append(note)
    elif note:
        sections.append(note)
    if gross < 0:
        sections.append('Der Betrag wird auf das bekannte Projektkonto der Auftraggeberin zurückgezahlt. Die Rechnungskorrektur ändert nur den vorstehend bezeichneten Leistungsumfang; die übrigen Positionen der ursprünglichen Rechnung bleiben bestehen.')
    else:
        due = (date.fromisoformat(dt) + timedelta(days=14)).isoformat()
        sections.append(f'Zahlbar ohne Skonto bis {date.fromisoformat(due).strftime("%d.%m.%Y")}. Bitte geben Sie bei der Überweisung ausschließlich die Rechnungsnummer {ident} und {REF} an. Zahlungsempfänger ist das im Briefkopf bezeichnete Unternehmen; die hinterlegte Bankverbindung bleibt unverändert. Die Leistung wurde für das Bauvorhaben und den angegebenen Zeitraum erbracht. Rückfragen zur Mengenaufstellung beantwortet die zeichnende Projektleitung.')
    rec(n, slug+'.pdf', dt, actor, f'{kind} {ident} – {label}', sections)
    return obj


def other(n, slug, dt, actor, ident, label, kg, amount, sections, *, kind='Bescheid', payment=None):
    INVOICES.append(dict(number=n, filename=f'{n:03d}_{slug}.pdf', date=dt, issuer=actor, id=ident,
                         label=label, kg=kg, net=amount, vat=0, gross=amount, rate=0,
                         period=dt, kind=kind, items=[], prior=None, payment=payment))
    rec(n, slug+'.pdf', dt, actor, label, sections)


def build():
    RECORDS.clear(); INVOICES.clear()
    other(70, 'Kaufpreisfaelligkeit_Grundstueck', '2026-11-16', 'notar', 'KP-2026-118', 'Kaufpreisfälligkeit Grundstück', '100', 450000, [
        'Zur Urkunde vom 03.11.2026, UVZ 118/2026, teile ich mit, dass die vereinbarten Fälligkeitsvoraussetzungen für den Kaufpreis eingetreten sind. Die Auflassungsvormerkung ist eingetragen und das kommunale Negativzeugnis liegt vor. Nicht übernommene Altbelastungen bestehen nicht; eine Löschung von Verkäufergrundpfandrechten ist nicht erforderlich. Das Grundstück ist das im Kaufvertrag bezeichnete Baugrundstück Wohnhof Am Steinbogen 18 in Hildesheim.',
        ('h', '1. Kaufpreis und Zahlung'),
        'Der Kaufpreis beträgt 450.000,00 EUR. Bitte überweisen Sie den Betrag bis 30.11.2026 unmittelbar an die Verkäuferin. Zahlungsempfängerin ist die Leinequartier Grundstücksverwaltung GmbH; es handelt sich nicht um eine Einzahlung auf ein Notaranderkonto. Die Bankverbindung ist unverändert die in der Urkunde angegebene Verbindung.',
        'Eine Umsatzsteueroption wurde im Grundstückskaufvertrag nicht erklärt. Neben dem Kaufpreis sind keine weiteren Kaufpreisbestandteile, Vermittlungsvergütungen oder beweglichen Gegenstände vereinbart. Die gesondert anfallenden Erwerbsnebenkosten werden von den jeweils zuständigen Stellen erhoben.',
        'Bitte übersenden Sie den Zahlungsbeleg an die Verkäuferin und an mein Büro. Die Eigentumsumschreibung wird nach Eingang der steuerlichen Unbedenklichkeitsbescheinigung und Nachweis der vollständigen Kaufpreiszahlung weiterbetrieben. Diese Mitteilung betrifft nur die vertraglich vereinbarten Fälligkeitsvoraussetzungen.'], kind='Kaufpreis', payment='2026-11-27')
    other(71, 'Grunderwerbsteuerbescheid', '2026-11-20', 'steuer', 'GRE-26-816', 'Grunderwerbsteuerbescheid', '100', 22500, [
        'Für den Erwerb aufgrund des notariellen Kaufvertrags vom 03.11.2026, UVZ 118/2026, wird gegenüber der Steinbogen Wohnen GmbH Grunderwerbsteuer in Höhe von 22.500,00 EUR festgesetzt. Gegenstand ist das Baugrundstück Wohnhof Am Steinbogen 18 in Hildesheim. Der Erwerbsvorgang wird unter der Steuernummer 30/145/60816 geführt.',
        ('h', '1. Festsetzung und Fälligkeit'), ('t', [['Berechnungsgrundlage', 'EUR'], ['Gegenleistung laut Kaufvertrag', '450.000,00'], ['Steuersatz Niedersachsen 5 %', '22.500,00'], ['Festgesetzte Steuer', '22.500,00']], [355,140]),
        'Der festgesetzte Betrag ist am 20.12.2026 fällig. Bitte verwenden Sie für die Zahlung das Kassenzeichen GRE-26-816 und die Ihnen bekannte Bankverbindung der Finanzkasse. Umsatzsteuer wird auf diese Steuerfestsetzung nicht erhoben.',
        'Die Festsetzung beruht auf der vorgelegten Erwerbsanzeige und dem Kaufvertrag. Weitere Gegenleistungen wurden nicht mitgeteilt. Die Unbedenklichkeitsbescheinigung wird nach Zahlung der Steuer an die beurkundende Notarin übersandt.',
        ('h', '2. Rechtsbehelfsbelehrung'), 'Gegen diesen Bescheid kann innerhalb eines Monats nach Bekanntgabe Einspruch beim Finanzamt Hildesheim eingelegt werden. Der Einspruch ist schriftlich, elektronisch oder zur Niederschrift zu erklären. Die Einlegung eines Einspruchs hemmt die Fälligkeit nicht.'], payment='2026-12-08')
    other(72, 'Notarkosten_Schlussberechnung', '2026-12-22', 'notar', 'NF-26-118-139', 'Notarkostenberechnung Grundstückskauf und Grundschuld', '100', 7850.51, [
        'Kostenberechnung zu UVZ 118/2026 vom 03.11.2026 und UVZ 139/2026 vom 19.12.2026. Kostenschuldnerin: Steinbogen Wohnen GmbH, Wohnhof Am Steinbogen 18, 31134 Hildesheim. Steuernummer der Notarin: 30/145/13705. Leistungszeitraum 03.11.2026 bis 22.12.2026.',
        'Der Geschäftswert des Kaufvertrags beträgt 450.000 EUR; die einfache Gebühr nach Tabelle B beträgt 885 EUR. Der Geschäftswert der Grundschuldbestellung mit Vollstreckungsunterwerfung beträgt 1.600.000 EUR; die einfache Gebühr beträgt 2.695 EUR. Eine Altbelastung der Verkäuferin war nicht zu löschen. Ein Anderkonto und zusätzliche Treuhandauflagen bestanden nicht.',
        ('t', [['KV GNotKG / Tätigkeit', 'Ansatz', 'Betrag EUR'],
               ['21100 Kaufvertrag', '2,0 × 885', '1.770,00'],
               ['22110 / 22112 Vollzug: Anforderung und Prüfung des kommunalen Negativzeugnisses', 'begrenzt', '50,00'],
               ['22200 Nr. 2 Prüfung und Mitteilung der Kaufpreisfälligkeit', '0,5 × 885', '442,50'],
               ['22115 strukturierte Daten Kaufvertrag', '0,1 × 885', '88,50'],
               ['21200 Grundschuldbestellung', '1,0 × 2.695', '2.695,00'],
               ['22200 Nr. 7 Herbeiführung der Bindung für die Gläubigerin durch gesonderte Entgegennahme', '0,5 × 2.695', '1.347,50'],
               ['22114 strukturierte Daten Grundschuld', 'Höchstbetrag', '125,00'],
               ['32001 abgegebene S/W-Seiten: Kauf 60 Seiten / Grundschuld 40 Seiten', '9 + 6', '15,00'],
               ['32005 Post- und Telekommunikationspauschalen für beide Urkunden', '2 × 20', '40,00'],
               ['32011 zwei Grundbuchabrufe, JVKostG KV 1151', '2 × 8', '16,00'],
               ['Steuerpflichtige Gebühren und Auslagen', '', '6.589,50'],
               ['32014 Umsatzsteuer 19 %', '', '1.252,01'],
               ['32015 zwei Gebühren Urkundenarchiv gemäß § 2 UA-GebS als durchlaufende Beträge', '2 × 4,50', '9,00'],
               ['Zahlbetrag', '', '7.850,51']], [305,85,105]),
        ('page',),
        'Die Bindungsherbeiführung zur Grundschuld erfolgte durch gesonderte Entgegennahme für die Gläubigerin am 22.12.2026. Die Kosten für das Urkundenarchiv wurden für Rechnung der Kostenschuldnerin verauslagt und sind nicht Teil der umsatzsteuerpflichtigen Bemessungsgrundlage. Vorschüsse wurden nicht vereinnahmt.',
        'Bitte zahlen Sie 7.850,51 EUR bis 05.01.2027 unter der Rechnungsnummer NF-26-118-139 auf das Kanzleikonto. Gerichtskosten werden gesondert von der Gerichtskasse erhoben. Gegen diese Kostenberechnung kann die Entscheidung des zuständigen Landgerichts nach § 127 GNotKG beantragt werden.'], kind='Kostenberechnung', payment='2027-01-05')
    INVOICES[-1].update(net=6598.50, vat=1252.01, taxable_base=6589.50, rate=.19)
    other(73, 'Grundbuch_Kostenrechnung', '2026-12-23', 'grundbuch', 'GB-26-778', 'Gerichtskostenrechnung Grundbuch', '100', 4047.50, [
        'In der Grundbuchsache Steinbogen Wohnen GmbH, Erwerb Wohnhof Am Steinbogen 18, Hildesheim, werden die nachstehenden Gerichtskosten angesetzt. Die Anträge sind unter dem Geschäftszeichen HI-GB-2026-778 erfasst. Kostenschuldnerin ist die Erwerberin nach den eingereichten Anträgen.',
        ('t', [['KV GNotKG / Eintragung', 'Geschäftswert EUR', 'Betrag EUR'], ['14110 Eigentumsumschreibung, 1,0', '450.000,00', '885,00'], ['14121 Buchgrundschuld, 1,0', '1.600.000,00', '2.695,00'], ['14150 Auflassungsvormerkung, 0,5', '450.000,00', '442,50'], ['14152 Löschung der Auflassungsvormerkung', 'Festgebühr', '25,00'], ['Gesamtbetrag', '', '4.047,50']], [290,110,95]),
        'Die Zahlung ist unter dem Kassenzeichen GB-26-778 bis 06.01.2027 an die Gerichtskasse zu leisten. Die mitgeteilte Bankverbindung der Gerichtskasse gilt unverändert. Es wurden keine Vorschüsse angerechnet. Umsatzsteuer wird auf die Gerichtsgebühren nicht erhoben.',
        'Der Ansatz enthält die Eigentumsumschreibung, die eingetragene Buchgrundschuld, die zuvor eingetragene Erwerbsvormerkung und deren Löschung. Gebühren einer Löschung fremder Altbelastungen sind nicht entstanden. Die Geschäftswerte wurden den Urkunden UVZ 118/2026 und 139/2026 entnommen.',
        'Gegen den Kostenansatz ist die Erinnerung nach § 81 GNotKG zulässig. Sie kann beim Amtsgericht Hildesheim schriftlich oder zur Niederschrift der Geschäftsstelle eingelegt werden. Bitte geben Sie bei Rückfragen Geschäfts- und Kassenzeichen an.'], payment='2027-01-06')
    inv(74, 'Maklerrechnung', '2026-11-04', 'makler', 'WG-26-104', 'Grundstücksvermittlung', '100', [('Nachweis und Vermittlung Baugrundstück gemäß Maklerauftrag vom 05.10.2026: 3 % des Kaufpreises von 450.000 EUR',1,'Auftrag',13500)], '05.10.2026 bis 03.11.2026', payment='2026-11-18')
    other(75, 'Baugenehmigung_Kostenbescheid', '2027-09-21', 'bauamt', 'BA-27-328-K', 'Kostenbescheid zum Bauantrag', '700', 4957.40, [
        'Für das vereinfachte Baugenehmigungsverfahren zum Bauantrag SW-HI-26-08 für den Neubau eines Wohngebäudes mit acht Wohnungen, Wohnhof Am Steinbogen 18, werden Gebühren von 4.957,40 EUR festgesetzt. Die Genehmigungsentscheidung wurde gesondert bekannt gegeben. Rechtsgrundlage sind die Niedersächsische Baugebührenordnung (NBauGO) und die nachstehend genannten Gebührenpositionen.',
        ('h', '1. Gebührenübersicht'),
        'Der Brutto-Rauminhalt beträgt nach den Bauvorlagen 2.710,00 m³. Bei einem Rohbauwert von 208,00 EUR je m³ beträgt der rechnerische Rohbauwert 563.680,00 EUR. Er wird gemäß § 3 Absatz 4 NBauGO auf 564.000,00 EUR aufgerundet. Der zugrunde gelegte Rohbauwert je m³ entspricht der zum 01.10.2026 veröffentlichten Festsetzung.',
        ('t', [['Leistung und Berechnung', 'Gebühr EUR'], ['Anlage 1 Nr. 1.1 Buchstabe a NBauGO: 1.128 angefangene Einheiten zu 500 EUR × 4,30 EUR', '4.850,40'], ['Anlage 1 Nr. 12.7 NBauGO: Nachforderung vom 02.06.2027, 60 Minuten', '107,00'], ['Gesamtbetrag', '4.957,40']], [375,120]),
        'Für die Prüfung der unvollständigen Bauvorlagen und das Erstellen der Nachforderung vom 02.06.2027 wurden 60 Minuten einer Sachbearbeiterin der Laufbahngruppe 2 unterhalb des zweiten Einstiegsamts erfasst. Nach § 6 Absatz 2 NBauGO in Verbindung mit § 4 Absatz 3 Nummer 2 KOVerm werden zwei angefangene halbe Stunden zu je 53,50 EUR berechnet.',
        'Zahlen Sie den Gesamtbetrag bis 21.10.2027 unter Angabe des Kassenzeichens BA-27-328-K an die Stadtkasse. Es handelt sich um öffentlich-rechtliche Gebühren; Umsatzsteuer wird nicht ausgewiesen. Der Kostenbescheid enthält keine Entgelte privater Prüfingenieure oder Fachplaner.',
        'Gegen diesen Kostenbescheid kann innerhalb eines Monats nach Bekanntgabe Widerspruch bei der Stadt Hildesheim erhoben werden. Bei öffentlichen Kosten hat der Widerspruch gemäß § 80 Absatz 2 Satz 1 Nummer 1 VwGO keine automatische aufschiebende Wirkung. Für Rückfragen sind Aktenzeichen und Kassenzeichen anzugeben. Dieser Kostenbescheid ersetzt nicht die Genehmigungsunterlagen.'], payment='2027-10-05')
    other(76, 'Bauversicherung_Beitragsrechnung', '2027-11-22', 'versicherung', 'BV-27-180', 'Beitragsrechnung Bauleistungs- und Bauherrenhaftpflichtversicherung', '700', 9000, [
        'Versicherungsschein BV-SW-2027-180: Versicherungsnehmerin Steinbogen Wohnen GmbH; versichertes Bauvorhaben Wohnhof Am Steinbogen 18, Hildesheim, acht Wohneinheiten. Der vereinbarte Schutz besteht vom 01.12.2027 bis 30.11.2028 für die im Versicherungsschein bezeichnete Bauleistung und die Bauherrenhaftpflicht.',
        ('h', '1. Einmalbeitrag'), ('t', [['Bestandteil', 'EUR'], ['Beitrag ohne Versicherungsteuer', '7.563,03'], ['Versicherungsteuer 19 %', '1.436,97'], ['Zahlbetrag', '9.000,00']], [355,140]),
        'Der Gesamtbeitrag von 9.000,00 EUR wird am 06.12.2027 fällig. Diese Beitragsrechnung weist Versicherungsteuer und keine Umsatzsteuer aus. Bitte verwenden Sie bei der Zahlung die Vertragsnummer BV-SW-2027-180. Der Beitrag umfasst die im Antrag vereinbarte Versicherungssumme; Erweiterungen des Bauumfangs sind dem Versicherer mitzuteilen.',
        'Der Versicherungsschutz richtet sich nach dem Versicherungsschein und den vereinbarten Bedingungen. Schäden sind unter Angabe des Schadentags und der ergriffenen Maßnahmen unverzüglich dem Vertragsservice anzuzeigen. Die Beitragserhebung ersetzt keine Deckungsentscheidung zu einem konkreten Schaden.'], payment='2027-12-06')
    inv(77, 'Bank_Finanzierungsentgelt', '2027-01-02', 'bank', 'LB-27-010', 'Vereinbartes Finanzierungsentgelt', '800', [('Individuell vereinbartes Entgelt für die Bereitstellung der Projektfinanzierung gemäß Darlehensvertrag Ziffer 3',1,'Vereinbarung',10000)], '02.01.2027', rate=0, note='Die Kreditgewährung und das ihr zugeordnete Finanzierungsentgelt werden umsatzsteuerfrei gemäß § 4 Nr. 8 Buchstabe a UStG abgerechnet. Es wird keine Umsatzsteuer geschuldet oder ausgewiesen.', payment='2027-01-16')
    inv(78, 'Bank_Zinsabrechnung_2027', '2027-12-31', 'bank', 'LB-Z-2027', 'Darlehenszinsen 2027', '800', [('Tranche 400.000 EUR vom 01.01. bis 31.12.2027 zu 3,6 % jährlich; 360/360 Tage',1,'Zeitraum',14400),('Tranche 400.000 EUR vom 01.07. bis 31.12.2027 zu 3,6 % jährlich; 180/360 Tage',1,'Zeitraum',7200),('Tranche 800.000 EUR vom 01.10. bis 31.12.2027 zu 3,6 % jährlich; 90/360 Tage',1,'Zeitraum',7200)], '01.01.2027 bis 31.12.2027', rate=0, note='Zinsen für die Kreditgewährung; steuerfrei gemäß § 4 Nr. 8 Buchstabe a UStG. Die Zinsen werden gesondert belastet und dem Darlehenskapital nicht zugeschlagen.', payment='2027-12-31')
    inv(79, 'Bank_Zinsabrechnung_2028', '2028-12-31', 'bank', 'LB-Z-2028', 'Darlehenszinsen 2028', '800', [('Darlehenssaldo 1.600.000 EUR vom 01.01. bis 31.12.2028 zu 3,6 % jährlich; 360/360 Tage',1,'Zeitraum',57600),('Vertragliche einmalige Zinsvergütung aus Darlehensvertrag Ziffer 3',1,'Vergütung',-6400)], '01.01.2028 bis 31.12.2028', rate=0, note='Zinsen für die Kreditgewährung; steuerfrei gemäß § 4 Nr. 8 Buchstabe a UStG. Die Zinsvergütung von 6.400 EUR ist eine Leistung der Bank aus der individuellen Finanzierungsvereinbarung und kein öffentlicher Zuschuss. Die Abrechnung umfasst auch die Zeit nach Bezugsbeginn; sie trifft keine Aussage über eine bilanzielle Aktivierung.', payment='2028-12-31')
    inv(80, 'Architekt_Rechnung_LPH1_4', '2027-09-24', 'architekt', 'KF-27-091', 'Objektplanung LPH 1 bis 4', '700', [('LPH 1 Grundlagenermittlung, vereinbarter Teilbetrag',1,'Phase',2900),('LPH 2 Vorplanung, vereinbarter Teilbetrag',1,'Phase',10150),('LPH 3 Entwurfsplanung, vereinbarter Teilbetrag',1,'Phase',21750),('LPH 4 Genehmigungsplanung, vereinbarter Teilbetrag',1,'Phase',4350)], '10.10.2026 bis 21.09.2027', note='Pauschalhonorar Objektplanung 145.000 EUR netto für die vereinbarten LPH 1 bis 9. Abgerechnet werden hier 27 % nach dem vertraglichen Zahlungsplan. Weitere Leistungsbilder und Nebenkosten werden nicht zusätzlich berechnet.', payment='2027-10-08')
    inv(81, 'Architekt_Rechnung_LPH5_7', '2027-11-26', 'architekt', 'KF-27-116', 'Objektplanung LPH 5 bis 7', '700', [('LPH 5 Ausführungsplanung',1,'Phase',36250),('LPH 6 Vorbereitung der Vergabe',1,'Phase',14500),('LPH 7 Mitwirkung bei der Vergabe',1,'Phase',5800)], '01.06.2027 bis 24.11.2027', note='Abgerechnet werden 39 % des vereinbarten Pauschalhonorars von 145.000 EUR netto. Die Teilbeträge für LPH 1 bis 4 wurden mit KF-27-091 abgerechnet; sie sind nicht erneut Bestandteil dieser Rechnung.', payment='2027-12-10')
    inv(82, 'Architekt_Rechnung_LPH8', '2028-10-20', 'architekt', 'KF-28-104', 'Objektplanung LPH 8', '700', [('Objektüberwachung und Dokumentation gemäß vereinbartem Leistungsumfang',1,'Phase',46400)], '01.12.2027 bis 15.10.2028', note='Abgerechnet werden 32 % des Pauschalhonorars von 145.000 EUR netto. Die Leistungen der Objektbetreuung in LPH 9 und deren Teilhonorar von 2.900 EUR netto bleiben außerhalb dieser Rechnung.', payment='2028-11-03')
    inv(83, 'Architekt_Schlussrechnung_LPH9', '2033-09-26', 'architekt', 'KF-33-092', 'Objektplanung LPH 9 und Schlussabrechnung', '700', [('Objektbetreuung gemäß Leistungsvereinbarung einschließlich Vorfristbegehung und Dokumentation',1,'Phase',2900)], '16.09.2028 bis 25.09.2033', prior=['KF-27-091','KF-27-116','KF-28-104'], note='Das vereinbarte Pauschalhonorar beträgt abschließend 145.000 EUR netto. Die bereits vereinnahmten Abschläge werden mit ihren Nettobeträgen und der darauf entfallenden Umsatzsteuer abgesetzt. Zusätzliche Mangelbeseitigungsüberwachung wird nicht mit dieser Rechnung abgerechnet.', payment='2033-09-30')
    inv(84, 'TGA_Planungsrechnung', '2027-09-27', 'tga', 'PT-27-092', 'TGA Planung', '700', [('Planung Wärmeversorgung und Wohnungsstationen',1,'Leistung',14000),('Planung Trinkwasser, Abwasser und Lüftung',1,'Leistung',10000),('Koordination Elektro, PV und Ladeinfrastruktur',1,'Leistung',6000)], '01.12.2026 bis 21.09.2027', payment='2027-10-11')
    inv(85, 'TGA_Schlussrechnung', '2028-10-22', 'tga', 'PT-28-104', 'TGA Ausführungsbegleitung und Abschluss', '700', [('Ausführungsdetails und gewerkeübergreifende Koordination',1,'Leistung',12000),('Mitwirkung Ausschreibung und Angebotsprüfung',1,'Leistung',8000),('Fachbauüberwachung, Prüfprotokolle und Übergabe',1,'Leistung',15000)], '22.09.2027 bis 15.10.2028', prior=['PT-27-092'], payment='2028-11-05')
    inv(86, 'Tragwerk_Schlussrechnung', '2028-09-22', 'tragwerk', 'BW-28-093', 'Tragwerksplanung', '700', [('Statische Berechnungen einschließlich Nachweisen',1,'Auftrag',18000),('Schal- und Bewehrungsplanung',1,'Auftrag',17000),('Bewehrungsabnahmen und Abschlussdokumentation',1,'Auftrag',5000)], '01.12.2026 bis 15.09.2028', payment='2028-10-06')
    inv(87, 'Vermessung_Schlussrechnung', '2028-09-24', 'vermessung', 'LP-28-094', 'Vermessungsleistungen', '700', [('Lage- und Höhenaufnahme',1,'Auftrag',6000),('Gebäudeabsteckung und Kontrollmessungen',1,'Auftrag',5000),('Schlussvermessung und Bestandsunterlagen',1,'Auftrag',4000)], '15.10.2026 bis 20.09.2028', payment='2028-10-08')
    inv(88, 'Baugrund_Rechnung', '2027-01-18', 'geo', 'BP-27-018', 'Baugrunduntersuchung', '700', [('Feldarbeiten mit sechs Sondierungen und Probenentnahmen',6,'Punkt',1100),('Laboruntersuchungen und Gründungsgutachten',1,'Auftrag',5400)], '02.11.2026 bis 13.01.2027', payment='2027-02-01')
    inv(89, 'Energie_Rechnung', '2028-09-25', 'energie', 'EP-28-096', 'Energieplanung und Fertigstellungsnachweis', '700', [('Energiebilanz und Wärmebrückennachweise',1,'Auftrag',6500),('Baubegleitung und energetische Abschlussdokumentation',1,'Auftrag',3500)], '15.11.2026 bis 20.09.2028', payment='2028-10-09')
    inv(90, 'Brandschutz_Rechnung', '2028-09-25', 'brand', 'RW-28-095', 'Brandschutzplanung', '700', [('Brandschutznachweis und Planabstimmung',1,'Auftrag',5500),('Baubegleitende Begehung und Schlussdokumentation',1,'Auftrag',2500)], '15.11.2026 bis 20.09.2028', payment='2028-10-09')
    inv(91, 'Rohbau_Abschlag1', '2028-02-15', 'rohbau', 'SH-28-021', 'Rohbau 1. Abschlag', '300', [('Baustelleneinrichtung und Vorhaltung gemäß Auftrag',1,'psch',20000),('Baugrubenaushub einschließlich vereinbarter Entsorgung',1000,'m³',45),('Gründung und Bodenplatte, Beton einschließlich Bewehrung',225,'m³',600)], '01.12.2027 bis 12.02.2028', payment='2028-02-29')
    inv(92, 'Rohbau_Abschlag2', '2028-04-15', 'rohbau', 'SH-28-046', 'Rohbau 2. Abschlag', '300', [('Gründung, verbleibender Beton einschließlich Bewehrung',25,'m³',600),('Tragende Wandkonstruktionen',1000,'m²',180),('Geschossdecken einschließlich Schalung und Bewehrung',262.5,'m²',400)], '13.02.2028 bis 12.04.2028', note='Der Abschlag erfasst ausschließlich die seit SH-28-021 hinzugekommenen und gemeinsam aufgemessenen Leistungen. Die Gründungsmengen aus beiden Abschlägen ergeben zusammen 250 m³. Der gesonderte Drännachtrag wird nicht in diese Grundvertragsabrechnung einbezogen.', payment='2028-04-29')
    inv(93, 'Rohbau_Schlussrechnung', '2028-09-30', 'rohbau', 'SH-28-093', 'Rohbau Schlussrechnung', '300', [('Geschossdecken, Restmenge',187.5,'m²',400),('Treppenläufe einschließlich Podesten',3,'Einheit',25000)], '13.04.2028 bis 15.09.2028', prior=['SH-28-021','SH-28-046'], note='Grundvertrag 650.000 EUR netto. Der gesondert beauftragte Nachtrag SH-N01 ist zusätzlich in SH-28-N01 abgerechnet und wird hier weder angesetzt noch abgesetzt.', payment='2028-10-14')
    inv(94, 'Gebaeudehuelle_Abschlag', '2028-05-15', 'huelle', 'DG-28-052', 'Gebäudehülle Abschlag', '300', [('Dachaufbau einschließlich Abdichtung und Wärmedämmung',400,'m²',300),('Fenstereinheiten gemäß Fensterliste, erster Abschnitt',50,'Stück',2000)], '01.04.2028 bis 12.05.2028', payment='2028-05-29')
    inv(95, 'Gebaeudehuelle_Schlussrechnung', '2028-09-30', 'huelle', 'DG-28-098', 'Gebäudehülle Schlussrechnung', '300', [('Fenstereinheiten gemäß Fensterliste, zweiter Abschnitt',50,'Stück',2000),('Fassadenbekleidung einschließlich Unterkonstruktion und Dämmung',800,'m²',150)], '13.05.2028 bis 15.09.2028', prior=['DG-28-052'], payment='2028-10-14')
    inv(96, 'HLS_Abschlag', '2028-06-15', 'hls', 'WH-28-063', 'Heizung Lüftung Sanitär Abschlag', '400', [('Wärmepumpenanlage einschließlich Speicher und Regelung',1,'Anlage',80000),('Rohrleitungen einschließlich Dämmung und Befestigung',1000,'m',60),('Fußbodenheizung beheizte Flächen',500,'m²',80)], '01.05.2028 bis 12.06.2028', payment='2028-06-29')
    inv(97, 'HLS_Schlussrechnung', '2028-09-30', 'hls', 'WH-28-094', 'Heizung Lüftung Sanitär Schlussrechnung', '400', [('Sanitärausstattung und Wohnungsanschlüsse',8,'Wohnung',12000),('Lüftungseinheiten und Verteilung',8,'Wohnung',8000),('Inbetriebnahme, Abgleich und Dokumentation',1,'Auftrag',20000)], '13.06.2028 bis 15.09.2028', prior=['WH-28-063'], payment='2028-10-14')
    inv(98, 'Elektro_Abschlag', '2028-06-01', 'elektro', 'LE-28-061', 'Elektro Abschlag', '400', [('Wohnungsinstallationen gemäß Raumbuch',8,'Wohnung',6500),('Hauptverteilung und Messkonzept',1,'Anlage',23000)], '01.05.2028 bis 30.05.2028', payment='2028-06-15')
    inv(99, 'Elektro_PV_Schlussrechnung', '2028-09-30', 'elektro', 'LE-28-098', 'Elektro und Ladeinfrastruktur Schlussrechnung', '400', [('Vorverkabelung für acht Stellplätze und zwei Ladepunkte gemäß Ausführungsplan',1,'Auftrag',18000),('Allgemeinbeleuchtung und Außeninstallation',1,'Auftrag',6000),('Aufzugszuleitung, Prüfungen und Dokumentation',1,'Auftrag',6000)], '31.05.2028 bis 15.09.2028', prior=['LE-28-061'], note='Der Elektroauftrag ohne die separat gelieferte PV-Anlage beträgt 105.000 EUR netto. Die eigenständige PV-Lieferung und Montage über 45.000 EUR wird ausschließlich mit LE-PV-28-099 abgerechnet. Ein Mieterstromliefermodell ist nicht Gegenstand dieses Auftrags. Die Zahlungseingänge zur Abschlagsrechnung werden außerhalb der Schlussabrechnung in der Debitorenbuchhaltung zugeordnet.', payment='2028-10-14')
    inv(100, 'Innenausbau_Abschlag', '2028-07-15', 'ausbau', 'IA-28-074', 'Innenausbau Abschlag', '300', [('Estrich einschließlich Vorbereitung und Randstreifen',1000,'m²',40),('Innenputz und Spachtelarbeiten',2200,'m²',50)], '01.06.2028 bis 12.07.2028', payment='2028-07-29')
    inv(101, 'Innenausbau_Schlussrechnung', '2028-09-30', 'ausbau', 'IA-28-095', 'Innenausbau Schlussrechnung', '300', [('Bodenbeläge gemäß Raumbuch',600,'m²',100),('Innentüren einschließlich Zargen und Beschlägen',80,'Stück',500),('Malerarbeiten einschließlich Untergrundbehandlung',2500,'m²',20)], '13.07.2028 bis 15.09.2028', prior=['IA-28-074'], payment='2028-10-14')
    inv(102, 'Aufzug_Schlussrechnung', '2028-09-30', 'aufzug', 'HA-28-093', 'Aufzugsanlage', '400', [('Aufzugsanlage einschließlich Steuerung und Kabine',1,'Anlage',42000),('Montage und Anschlussarbeiten',1,'Auftrag',12000),('Prüfungen, Inbetriebnahme und Betreiberunterlagen',1,'Auftrag',6000)], '01.07.2028 bis 15.09.2028', payment='2028-10-14')
    inv(103, 'Aussenanlagen_Schlussrechnung', '2028-11-15', 'aussen', 'GG-28-115', 'Außenanlagen und Entwässerung', '500', [('Befestigte Flächen einschließlich Unterbau',500,'m²',100),('Außenentwässerung einschließlich Schächten',200,'m',200),('Pflanzflächen einschließlich Bodenvorbereitung',300,'m²',60),('Fahrradabstellanlage und Einfriedung',1,'Auftrag',12000)], '01.08.2028 bis 10.11.2028', payment='2028-11-29')
    inv(104, 'Hausanschluesse_Rechnung', '2028-08-20', 'versorger', 'LN-28-081', 'Hausanschlüsse', '200', [('Trinkwasseranschluss Grundstück bis Übergabestelle',1,'Anschluss',8500),('Schmutzwasseranschluss und Übergabeschacht',1,'Anschluss',12000),('Stromanschluss einschließlich Netzanschlussschrank',1,'Anschluss',14500)], '01.07.2028 bis 15.08.2028', payment='2028-09-03')
    inv(105, 'Rohbau_Nachtrag_Draenpackung', '2028-03-05', 'rohbau', 'SH-28-N01', 'Nachtrag Drän- und Filterpackung', '300', [('Zusätzlicher Aushub gemäß Freigabe vom 05.02.2028',60,'m³',85),('Filtermaterial liefern und einbauen',45,'m³',160),('Zusätzliche Dränleitung mit Filtervlies',30,'m',95),('Zusätzliche Arbeitsstunden gemäß bestätigtem Stundenbericht',57,'Stunde',50)], '07.02.2028 bis 11.02.2028', note='Grundlage ist die schriftliche Beauftragung SH-N01 vom 05.02.2028 nach dem Ortstermin zur abweichenden Bodenschicht. Der Preis von 18.000 EUR netto ergänzt den Grundvertrag. Die Mengen und Stunden wurden im gemeinsamen Aufmaß vom 11.02.2028 bestätigt.', payment='2028-03-19')
    inv(106, 'Innenausbau_Rechnungskorrektur', '2028-10-15', 'ausbau', 'IA-G-28-017', 'Rechnungskorrektur zu IA-28-095', '300', [('Vereinbarte Entgeltminderung für die abweichende Belagsausführung in Wohnung 07 gemäß Vereinbarung vom 15.10.2028',1,'Vereinbarung',-2500)], '15.10.2028', kind='Rechnungskorrektur', note='Diese Korrektur bezieht sich ausschließlich auf IA-28-095 vom 30.09.2028. Sie ist eine Entgeltminderung des leistenden Unternehmers und keine Abrechnungsgutschrift des Auftraggebers.', payment='2028-11-15')
    inv(119, 'PV_Lieferung_Montage_Nullsteuersatz', '2028-09-30', 'elektro', 'LE-PV-28-099', 'PV-Anlage Lieferung und Montage', '400', [('Solarmodule mit je 450 W Nennleistung einschließlich Unterkonstruktion',80,'Modul',400),('Wechselrichter, Anlagenverkabelung und Inbetriebnahme',1,'Anlage',13000)], '01.08.2028 bis 15.09.2028', rate=0, note='Umsatzsteuersatz 0 % gemäß § 12 Absatz 3 Nummern 1 und 4 UStG. Lieferung und Installation erfolgen an die Steinbogen Wohnen GmbH als Betreiberin der 36-kWp-Anlage auf dem Dach ihres Wohngebäudes mit acht Wohnungen. Gebäudeart und Betreiberstellung sind im Auftrag und in den Übergabeunterlagen dokumentiert. Die Leistung betrifft die PV-Anlage; allgemeine Elektroinstallation und Ladepunkte sind nicht in dieser Rechnung enthalten.', payment='2028-10-14')

    rec(107, 'Darlehensvereinbarung.docx', '2026-12-19', 'bank', 'Projektfinanzierung SW-HI-26-08', [
        ('h','1. Vertragsparteien und Kredit'), 'Die Leinebogen Projektbank AG gewährt der Steinbogen Wohnen GmbH zur Finanzierung des Neubaus Wohnhof Am Steinbogen 18 in Hildesheim ein Darlehen von 1.600.000,00 EUR. Die Geschäftsführerin Maren Birk handelt für die Darlehensnehmerin. Das Darlehen wird auf das ausschließlich für die Bauinvestition geführte Projektkonto ausgezahlt.',
        ('h','2. Auszahlung und Besicherung'), 'Die Bank zahlt 400.000,00 EUR zum 01.01.2027, 400.000,00 EUR zum 01.07.2027 und 800.000,00 EUR zum 01.10.2027 aus. Voraussetzung sind der Nachweis des Grundstückserwerbs und die vereinbarte erstrangige Grundschuld von 1.600.000,00 EUR. Die Sicherungszweckerklärung wird gesondert beurkundet beziehungsweise vereinbart. Eine weitergehende Sicherheit für fremde Verbindlichkeiten wird mit diesem Vertrag nicht bestellt.',
        ('h','3. Zinsen und Entgelt'), 'Der feste Sollzinssatz beträgt 3,60 % jährlich. Für die Bauphase bis 31.12.2028 werden die Zinsen nach der kaufmännischen Methode 30/360 berechnet und jeweils zum Jahresende dem Projektkonto belastet. Zusätzlich ist ein zwischen den Parteien individuell ausgehandeltes Finanzierungsentgelt von 10.000,00 EUR bei der ersten Auszahlung fällig. Die Bank gewährt für das Projekt einmalig eine Zinsvergütung von 6.400,00 EUR, die in der Zinsabrechnung 2028 verrechnet wird. Sie ist von keiner öffentlichen Förderzusage abhängig.',
        ('page',), ('h','4. Tilgung und Zahlungsweg'), 'Bis 31.12.2028 ist das Darlehen tilgungsfrei. Ab 01.01.2029 beträgt die anfängliche jährliche Tilgung 1,50 % des ursprünglichen Darlehensbetrags. Zusammen mit dem Sollzins ergibt sich eine jährliche Annuität von 81.600,00 EUR, zahlbar in zwölf gleichen Monatsraten von 6.800,00 EUR vom gesonderten Mietbetriebskonto. Innerhalb der gleichbleibenden Rate wächst der Tilgungsanteil, soweit der Zinsanteil durch die Tilgung sinkt. Der Sollzinssatz ist bis 31.12.2038 gebunden. Die Anschlussregelung und die dann verbleibende Restschuld werden vor Ablauf gesondert vereinbart.',
        ('h','5. Projektkonto und Unterlagen'), 'Das Projektkonto wird bis zum Ausgleich der letzten Architektenforderung ohne Habenzins und ohne weitere Kontoführungsentgelte geführt. Mieten, Mietkautionen und laufende Bewirtschaftungsvorgänge werden über andere Konten abgewickelt. Die Darlehensnehmerin übersendet der Bank die Kostenfortschreibung und wesentliche Änderungen des Bauvorhabens. Die Bank darf Auszahlungen nur bei Fehlen einer vereinbarten Auszahlungsvoraussetzung zurückhalten.',
        ('h','6. Abschluss'), 'Die Parteien haben die Bestimmungen gemeinsam am 19.12.2026 besprochen. Vertragsänderungen werden in Textform festgehalten. Gesetzliche Kündigungsrechte bleiben unberührt. Beide Parteien erhalten eine gleichlautende Ausfertigung. Für die Bank: Petra West. Für die Darlehensnehmerin: Maren Birk. Beide haben die Vereinbarung am 19.12.2026 unterzeichnet.'])
    rec(108, 'Eigenkapital_und_Vermietungsannahmen.docx', '2026-12-20', 'bauherr', 'Gesellschafterbeschluss zur Projektfinanzierung und Vermietungsplanung', [
        ('h','1. Eigenkapital'), 'Die Alleingesellschafterin Steinbogen Beteiligungen GmbH beschließt, der Steinbogen Wohnen GmbH für das Projekt SW-HI-26-08 insgesamt 2.100.000,00 EUR als freiwillige Zuzahlung in die Kapitalrücklage zur Verfügung zu stellen. Die Einzahlung erfolgt in drei Teilbeträgen von 800.000,00 EUR zum 15.11.2026, 700.000,00 EUR zum 15.01.2027 und 600.000,00 EUR zum 15.01.2028. Der erste Teilbetrag ist bereits auf dem Projektkonto eingegangen. Eine laufende Verzinsung oder Rückzahlungsfälligkeit wird nicht vereinbart.',
        ('h','2. Freigegebener Kostenrahmen'), 'Der freigegebene Kostenrahmen beträgt 3.650.000,00 EUR einschließlich nicht abziehbarer Umsatzsteuer, Grundstück und Erwerbsnebenkosten sowie der gesondert vereinbarten Finanzierungskosten bis Ende 2028. Die Ausgangskalkulation beträgt 3.463.770,41 EUR. Aus Eigenkapital und Darlehen stehen 3.700.000,00 EUR zur Verfügung. Nicht verbrauchte Mittel bleiben zunächst zur Deckung ausstehender Projektabrechnungen auf dem Projektkonto.',
        ('page',), ('h','3. Vermietungsannahmen'), 'Geplant ist die langfristige Vermietung sämtlicher acht Wohnungen zu Wohnzwecken ab 01.10.2028. Die Wohnflächen betragen in der Reihenfolge WE 01 bis WE 08 jeweils 68, 75, 82, 68, 75, 82, 75 und 75 m², zusammen 600 m². Für die Wirtschaftlichkeitsrechnung setzt die Geschäftsführung eine Nettokaltmiete von 13,50 EUR je m² und Monat an. Acht Stellplätze werden ausschließlich zusammen mit den Wohnungen mit jeweils 50,00 EUR monatlich kalkuliert. Diese Werte sind interne Planannahmen und keine Feststellung einer ortsüblichen oder rechtlich zulässigen Miethöhe.',
        ('h','4. Betrieb und Rücklagenplanung'), 'Für nicht umlagefähige Bewirtschaftungsaufwendungen und laufende Instandhaltung werden in einem stabilisierten vollen Vermietungsjahr zunächst 18.000,00 EUR veranschlagt. Diese Planannahme ist vor der Vermietung anhand von Angeboten und den tatsächlichen Verträgen zu aktualisieren. Mietausfall wird zunächst mit null angesetzt; die Arbeitsmappe soll abweichende Leerstands- und Kostenannahmen rechnen können. Kautionen stehen nicht zur Finanzierung der Baukosten zur Verfügung.',
        ('h','5. Nutzung und Verantwortlichkeit'), 'Die Gesellschaft erbringt keine Bauleistungen für Dritte. Sie baut und hält das Haus ausschließlich zur eigenen Wohnraumvermietung. Der PV-Ertrag wird im vereinbarten Allgemeinstromkonzept genutzt; gesonderte Stromlieferverträge mit Mietern werden nicht abgeschlossen. Die steuerliche Zuordnung und die Feststellung der Anschaffungs- und Herstellungskosten werden anhand der Belege mit der steuerlichen Beratung abgestimmt. Die kaufmännische Projektkostenrechnung stellt dafür die Zahlungen und Leistungsabrechnungen getrennt bereit.'])

    tx = []
    def t(dt, reference, counterparty, ident, amount, direction):
        tx.append(dict(date=dt, reference=reference, counterparty=counterparty, invoice=ident,
                       amount=round(amount,2), direction=direction))
    for dt, amount in [('2026-11-15',800000),('2027-01-15',700000),('2028-01-15',600000)]:
        t(dt,'EK-'+dt.replace('-',''), 'Steinbogen Beteiligungen GmbH','EK',amount,'Zufluss')
    for dt, amount in [('2027-01-01',400000),('2027-07-01',400000),('2027-10-01',800000)]:
        t(dt,'DL-'+dt.replace('-',''),ACTORS['bank'][0],'Darlehen',amount,'Zufluss')
    for i in INVOICES:
        payee='Leinequartier Grundstücksverwaltung GmbH' if i['number']==70 else ACTORS[i['issuer']][0]
        t(i['payment'], 'SW-'+str(i['number'])+'-'+i['payment'].replace('-',''), payee,i['id'], abs(i['gross']), 'Abfluss' if i['gross']>0 else 'Erstattung')
    t('2028-06-19','SW-98-20280619',ACTORS['elektro'][0],'LE-28-061',89250,'Abfluss')
    t('2028-12-15','LE-RET-20281215',ACTORS['elektro'][0],'LE-28-061',89250,'Erstattung')
    tx.sort(key=lambda x:(x['date'],x['reference']))
    balance=0
    for row in tx:
        balance += row['amount'] if row['direction'] != 'Abfluss' else -row['amount']
        row['balance'] = round(balance,2)
    for n, year in [(109,2026),(110,2027),(111,2028),(112,2033)]:
        subset=[x for x in tx if x['date'].startswith(str(year))]
        opening=round(subset[0]['balance'] - (subset[0]['amount'] if subset[0]['direction']!='Abfluss' else -subset[0]['amount']),2)
        rows=[['Valuta / Referenz','Buchungstext','Eingang EUR','Ausgang EUR','Saldo EUR']]
        for x in subset:
            rows.append([x['date']+'\n'+x['reference'],x['counterparty']+'\n'+x['invoice'],euro(x['amount']) if x['direction']!='Abfluss' else '–',euro(x['amount']) if x['direction']=='Abfluss' else '–',euro(x['balance'])])
        end_date='2033-09-30' if year==2033 else f'{year}-12-31'
        end_label='30.09.2033' if year==2033 else f'31.12.{year}'
        tables=[('t',rows,[111,165,73,73,73])]
        if year==2027:
            tables=[('t',rows[:9],[111,165,73,73,73]),
                    f'Übertrag: {euro(subset[7]["balance"])} EUR. Fortsetzung des Auszugs auf der nächsten Seite.',
                    ('page',),('h','Projektkonto 882608 – Fortsetzung 2027'),
                    ('t',[rows[0]]+rows[9:],[111,165,73,73,73])]
        rec(n,f'Projektkonto_{year}.pdf',end_date,'bank',f'Projektkonto 882608 – Auszug bis {end_label}',[
            f'Kontoinhaberin: Steinbogen Wohnen GmbH. Kontonummer 882608, Buchungskreis SW-HI-26-08. Auszugszeitraum 01.01.{year} bis {end_label}. Anfangssaldo {euro(opening)} EUR. Es werden sämtliche Umsätze dieses Projektkontos für den Auszugszeitraum aufgeführt.',
            *tables,
            f'Schlusssaldo {euro(subset[-1]["balance"])} EUR. Eingänge und Ausgänge sind mit ihrem jeweiligen Vorzeichen in den getrennten Spalten dargestellt. Der Saldo enthält ausschließlich Bewegungen des Projektkontos. Mietbetrieb, laufende Darlehenstilgung ab 2029 und Kautionskonten sind nicht Bestandteil dieses Kontoauszugs.',
            'Für 2029 bis 2032 bestanden auf diesem Projektkonto keine Umsätze. Die Mietbewirtschaftung wird seit Vermietungsbeginn auf dem gesonderten Betriebskonto geführt.' if year==2033 else 'Bitte reichen Sie Rückfragen zur Zuordnung einer Buchung unter Angabe von Auszugsjahr, Valuta und Bankreferenz ein. Eine auf dem Kontoauszug sichtbare Zahlung bestätigt keine fachliche Prüfung der zugrunde liegenden Rechnung.'])
    rec(113,'Rueckfrage_Elektrozahlung.eml','2028-09-26','bauherr','Zahlungszuordnung LE-28-061',[
        'Bei der Vorbereitung der Schlusszahlungen sehe ich im Projektkonto zwei Überweisungen mit dem Verwendungszweck LE-28-061, am 15.06. und am 19.06.2028, jeweils über 89.250 EUR. In meiner Liste steht nur die Rechnung vom 01.06.2028. Bitte senden Sie mir Ihr Debitorenkonto für das Bauvorhaben, damit wir die Zuordnung vor der Schlusszahlung abstimmen können.',
        'Die Schlussrechnungsunterlagen gehen parallel an Frau Feld zur technischen Prüfung. Bitte verrechnen Sie ohne kurze Abstimmung keine Beträge mit anderen Bauvorhaben. Die Überweisungsfreigabe am 19.06. wurde während meiner Abwesenheit durch die Vertretung vorgenommen; auf dem Freigabebeleg steht ebenfalls LE-28-061.',
        'Für die Bankbesprechung am 05.10. brauche ich eine belastbare Übersicht über den noch zu zahlenden Betrag. Ich sende Ihnen die Abschlagsrechnung nochmals mit. Bitte antworten Sie an die Projektadresse und setzen Sie die Buchhaltung in Ihrem Haus hinzu.'],recipient='elektro',attachments=['098_Elektro_Abschlag.pdf'])
    rec(114,'Elektro_Rueckzahlung.eml','2028-12-13','elektro','Debitorenkonto SW-HI-26-08 und Auszahlung',[
        'Wir haben die beiden Zahlungseingänge vom 15.06. und 19.06.2028 zu LE-28-061 nun auf demselben Debitorenkonto zusammengeführt. Der zweite Eingang war zunächst ohne Rechnungszuordnung verbucht. Die Schlussrechnung LE-28-098 wurde am 14.10.2028 in voller Höhe ausgeglichen.',
        'Auf dem Debitorenkonto verbleibt deshalb ein Guthaben von 89.250 EUR. Wir haben die Rückzahlung auf das Projektkonto 882608 für den 15.12.2028 angewiesen. Der Zahlungstext lautet LE-RET-20281215. Unsere Leistungsabrechnung und die ausgewiesene Umsatzsteuer ändern sich dadurch nicht; es handelt sich um die Rückzahlung einer zusätzlich eingegangenen Zahlung.',
        'Bitte teilen Sie uns mit, falls die Bankverbindung inzwischen geändert wurde. Andernfalls ist keine weitere Mitwirkung erforderlich. Die Abrechnung der Ladeinfrastruktur und der beiden Ladepunkte bleibt Bestandteil der bereits übersandten Schlussrechnung.'],attachments=['099_Elektro_PV_Schlussrechnung.pdf'])
    rec(115,'Innenausbau_Belagsvereinbarung.eml','2028-10-15','ausbau','Wohnung 07 – Entgeltminderung Bodenbelag',[
        'Wir bestätigen die heute mit Frau Birk vereinbarte Entgeltminderung von 2.500 EUR netto für die abweichende Belagsausführung in Wohnung 07. Die Auftraggeberin behält den eingebauten Belag; ein Austausch wird insoweit nicht mehr verlangt. Die sonstigen Ansprüche aus den übrigen Gewerkeleistungen werden durch diese Vereinbarung nicht geregelt.',
        'Die beigefügte Rechnungskorrektur IA-G-28-017 bezieht sich auf die Schlussrechnung IA-28-095. Da deren Betrag am 14.10. eingegangen ist, zahlen wir 2.975 EUR einschließlich Umsatzsteuer am 15.11.2028 auf das Projektkonto zurück. Eine Verrechnung mit anderen Objekten ist nicht vorgesehen.',
        'Bitte nehmen Sie diese Nachricht zusammen mit der Rechnungskorrektur zur Kostenfortschreibung. Wir übersenden keine neue Schlussrechnung, damit die ursprüngliche Rechnung und die einzelne Änderung anhand ihrer Nummern nachvollziehbar bleiben.'], attachments=['106_Innenausbau_Rechnungskorrektur.pdf'])
    rec(116,'Projektabschluss_Zahlungsstand.eml','2033-10-05','bauherr','Letzte Architektenzahlung und Unterlagen für die Jahresarbeit',[
        'Die Schlussrechnung KF-33-092 wurde am 30.09.2033 mit 3.451 EUR vom Projektkonto bezahlt. Bitte ordnen Sie die Rechnung der Objektbetreuung zu und geben Sie die Informationen an die steuerliche Beratung weiter. Die älteren Abschläge sind in der kumulierten Schlussrechnung abgesetzt; sie wurden nicht nochmals überwiesen.',
        'Die drei Projektarbeitsmappen enthalten die Belegbeträge, die Zahlungsliste des Projektkontos und die ursprünglichen Vermietungsannahmen. Bitte trennen Sie bei der Jahresarbeit die Investitionsabrechnung von den laufenden Ergebnissen der Vermietung. Das Betriebskonto für die Mietzahlungen, Nebenkosten und Annuitäten wird in der laufenden Buchhaltung geführt und gehört nicht zum Projektkonto.',
        'Die Notar- und Gerichtskosten liegen mit den Kostenberechnungen und den jeweiligen Zahlungen vor. Bitte berücksichtigen Sie bei der Belegzuordnung die gesondert ausgewiesenen durchlaufenden Archivkosten und die vom Grundstück getrennten Finanzierungsleistungen. Die beigefügte Architektenrechnung kann nun ebenfalls dem Zahlbeleg zugeordnet werden.'],attachments=['083_Architekt_Schlussrechnung_LPH9.pdf'])
    for number, filename, title in [(120,'120_Kostenentwicklung.xlsx','Kostenentwicklung'),(121,'121_Rechnungen_und_Zahlungen.xlsx','Rechnungen und Zahlungen'),(122,'122_Finanzierung_und_Vermietung.xlsx','Finanzierung und Vermietung')]:
        RECORDS.append(dict(number=number,filename=filename,date='2033-10-05',issuer='bauherr',recipient='bauherr',title=title,sections=[],attachments=[]))
    total=round(sum(i['gross'] for i in INVOICES),2)
    assert total==3482215.41, total
    assert round(tx[-1]['balance'],2)==217784.59
    groups = [
        ['100','Grundstück und Erwerbsnebenkosten',500463.01],['200','Hausanschlüsse',41650],
        ['300','Baukonstruktion',1654100],['400','Technische Anlagen',669750],['500','Außenanlagen',142800],
        ['700','Planung, Versicherung, Gebühren',365007.40],['800','Finanzierung bis Ende 2028',90000]]
    assert round(sum(x[2] for x in groups),2)==3463770.41
    create_xrechnungen()
    dataset=dict(root='.',case=str(CASE.relative_to(ROOT)),qa=str(QA),ref=REF,records=RECORDS,invoices=INVOICES,
                 transactions=tx,budget=3650000,base=3463770.41,total=total,groups=groups,
                 equity=2100000,loan=1600000,interest=.036,repayment=.015,
                 areas=[68,75,82,68,75,82,75,75],rent_sqm=13.5,parking_count=8,parking_rent=50,
                 operating_cost=18000,asof='2033-10-05',vat_deductible=0,
                 accounting_note='Projektmittelrechnung. Keine abschließende Feststellung handels- oder steuerrechtlicher Herstellungskosten.')
    DATA.parent.mkdir(parents=True,exist_ok=True); QA.mkdir(parents=True,exist_ok=True)
    DATA.write_text(json.dumps(dataset,ensure_ascii=False,indent=2)+'\n')
    build_records(RECORDS)
    (QA/'records.json').write_text(json.dumps(RECORDS,ensure_ascii=False,indent=2)+'\n')
    print(f'{len(INVOICES)} Kostenbelege, {len(tx)} Kontoumsätze, {len(RECORDS)} Originale; {euro(total)} EUR')
    return dataset


def create_xrechnungen():
    """UBL 2.1 gemäß XRechnung 3.0; Lesefassungen und XML sind derselbe Beleg."""
    from decimal import Decimal, ROUND_HALF_UP
    ns_invoice='urn:oasis:names:specification:ubl:schema:xsd:Invoice-2'
    ns_credit='urn:oasis:names:specification:ubl:schema:xsd:CreditNote-2'
    cac='urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2'
    cbc='urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2'
    ET.register_namespace('cac',cac); ET.register_namespace('cbc',cbc)
    def el(parent,tag,text=None,**attrs):
        namespace,name=tag.split(':');obj=ET.SubElement(parent,'{'+({'cac':cac,'cbc':cbc}[namespace])+'}'+name,attrs)
        if text is not None:obj.text=str(text)
        return obj
    def val(n):return str(Decimal(str(n)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
    def amount(parent,tag,n):return el(parent,'cbc:'+tag,val(n),currencyID='EUR')
    def iban(account):
        bban='00000000'+str(account).zfill(10);check=98-int(bban+'131400')%97
        return f'DE{check:02d}'+bban
    def party(parent,actor,seller=False):
        name,address,email,contact=ACTORS[actor];street,city=address.split(', ',1);postal,city=city.split(' ',1)
        p=el(parent,'cac:Party');el(p,'cbc:EndpointID',email,schemeID='EM')
        if seller:
            ident=el(p,'cac:PartyIdentification');el(ident,'cbc:ID','SW-KR-'+actor.upper())
        pn=el(p,'cac:PartyName');el(pn,'cbc:Name',name)
        a=el(p,'cac:PostalAddress');el(a,'cbc:StreetName',street);el(a,'cbc:CityName',city);el(a,'cbc:PostalZone',postal)
        co=el(a,'cac:Country');el(co,'cbc:IdentificationCode','DE')
        if seller:
            tax=el(p,'cac:PartyTaxScheme');el(tax,'cbc:CompanyID','30/145/'+str(12000+list(ACTORS).index(actor)*113))
            sch=el(tax,'cac:TaxScheme');el(sch,'cbc:ID','FC')
        legal=el(p,'cac:PartyLegalEntity');el(legal,'cbc:RegistrationName',name)
        if seller:
            c=el(p,'cac:Contact');el(c,'cbc:Name',contact);el(c,'cbc:Telephone','+49 5121 000000');el(c,'cbc:ElectronicMail',email)
    units={'m³':'MTQ','m²':'MTK','m':'MTR','Stunde':'HUR'}
    for invoice in INVOICES:
        if invoice['number']<80 or invoice['issuer']=='bank':continue
        if not invoice['items']:continue
        credit=invoice['gross']<0
        namespace=ns_credit if credit else ns_invoice
        ET.register_namespace('ubl',namespace)
        root=ET.Element('{'+namespace+'}'+('CreditNote' if credit else 'Invoice'))
        el(root,'cbc:CustomizationID','urn:cen.eu:en16931:2017#compliant#urn:xeinkauf.de:kosit:xrechnung_3.0')
        el(root,'cbc:ProfileID','urn:fdc:peppol.eu:2017:poacc:billing:01:1.0')
        el(root,'cbc:ID',invoice['id']);el(root,'cbc:IssueDate',invoice['date'])
        if not credit:el(root,'cbc:DueDate',(date.fromisoformat(invoice['date'])+timedelta(days=14)).isoformat())
        el(root,'cbc:CreditNoteTypeCode' if credit else 'cbc:InvoiceTypeCode','381' if credit else '380')
        el(root,'cbc:Note',invoice['label']+'. Bauvorhaben Wohnhof Am Steinbogen 18, Hildesheim. '+
           ('Rechnungskorrektur des Leistenden zu IA-28-095; Entgeltminderung Wohnung 07.' if credit else 'Die gleich nummerierte PDF-Datei ist die Lesefassung dieser Rechnung.'))
        el(root,'cbc:DocumentCurrencyCode','EUR');el(root,'cbc:BuyerReference',REF)
        dates=re.findall(r'\d{2}\.\d{2}\.\d{4}',invoice['period'])
        if dates:
            per=el(root,'cac:InvoicePeriod')
            def iso(t):return '-'.join(reversed(t.split('.')))
            el(per,'cbc:StartDate',iso(dates[0]));el(per,'cbc:EndDate',iso(dates[-1]))
        refs=['IA-28-095'] if credit else (invoice['prior'] or [])
        for ident in refs:
            earlier=next(x for x in INVOICES if x['id']==ident)
            br=el(root,'cac:BillingReference');ir=el(br,'cac:InvoiceDocumentReference');el(ir,'cbc:ID',ident);el(ir,'cbc:IssueDate',earlier['date'])
        sup=el(root,'cac:AccountingSupplierParty');party(sup,invoice['issuer'],True)
        customer=el(root,'cac:AccountingCustomerParty');party(customer,'bauherr')
        pay=el(root,'cac:PaymentMeans');el(pay,'cbc:PaymentMeansCode','58');el(pay,'cbc:PaymentID',invoice['id'])
        account=el(pay,'cac:PayeeFinancialAccount')
        el(account,'cbc:ID',iban(882608 if credit else 1000000000+list(ACTORS).index(invoice['issuer'])))
        el(account,'cbc:Name',ACTORS['bauherr' if credit else invoice['issuer']][0])
        pt=el(root,'cac:PaymentTerms');el(pt,'cbc:Note','Erstattung auf das Projektkonto.' if credit else 'Zahlbar ohne Skonto binnen 14 Tagen. Bereits vereinnahmte Abschläge sind abgesetzt.')
        items=[];prepaid=0
        for ident in (invoice['prior'] or []):
            earlier=next(x for x in INVOICES if x['id']==ident);items.extend(earlier['items']);prepaid+=earlier['gross']
        items.extend(invoice['items'])
        if credit:items=[(label,q,unit,abs(price)) for label,q,unit,price in items]
        net=round(sum(q*price for _,q,_,price in items),2);vat=float(Decimal(str(net*invoice['rate'])).quantize(Decimal('.01'),rounding=ROUND_HALF_UP));gross=round(net+vat,2)
        tax=el(root,'cac:TaxTotal');amount(tax,'TaxAmount',vat)
        subtotal=el(tax,'cac:TaxSubtotal');amount(subtotal,'TaxableAmount',net);amount(subtotal,'TaxAmount',vat)
        category=el(subtotal,'cac:TaxCategory');el(category,'cbc:ID','S' if invoice['rate'] else 'Z');el(category,'cbc:Percent',val(invoice['rate']*100))
        scheme=el(category,'cac:TaxScheme');el(scheme,'cbc:ID','VAT')
        monetary=el(root,'cac:LegalMonetaryTotal');amount(monetary,'LineExtensionAmount',net);amount(monetary,'TaxExclusiveAmount',net);amount(monetary,'TaxInclusiveAmount',gross)
        if prepaid:amount(monetary,'PrepaidAmount',prepaid)
        amount(monetary,'PayableAmount',round(gross-prepaid,2))
        for idx,(label,quantity,unit,price) in enumerate(items,1):
            line=el(root,'cac:CreditNoteLine' if credit else 'cac:InvoiceLine');el(line,'cbc:ID',str(idx))
            el(line,'cbc:CreditedQuantity' if credit else 'cbc:InvoicedQuantity',val(quantity),unitCode=units.get(unit,'C62'))
            amount(line,'LineExtensionAmount',quantity*price)
            item=el(line,'cac:Item');el(item,'cbc:Name',label)
            ct=el(item,'cac:ClassifiedTaxCategory');el(ct,'cbc:ID','S' if invoice['rate'] else 'Z');el(ct,'cbc:Percent',val(invoice['rate']*100));sch=el(ct,'cac:TaxScheme');el(sch,'cbc:ID','VAT')
            p=el(line,'cac:Price');amount(p,'PriceAmount',price)
        filename=Path(invoice['filename']).stem+'_XRechnung.xml'
        ET.indent(root,space='  ');(CASE/filename).write_bytes(ET.tostring(root,encoding='utf-8',xml_declaration=True))
        RECORDS.append(dict(number=invoice['number'],filename=filename,date=invoice['date'],issuer=invoice['issuer'],recipient='bauherr',title=invoice['label']+' – XRechnung',sections=[],attachments=[]))


def verify_workbooks():
    """Native Rechenprobe; OOXML nur für Drucklayout, Metadaten und Cache-Löschung."""
    from openpyxl import load_workbook
    from office_process import run_office
    reports=json.loads((QA/'workbook-qa.json').read_text())
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    def decorate(file, sheets, clear=False):
        output=io.BytesIO()
        with zipfile.ZipFile(file) as source, zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as dest:
            for item in source.infolist():
                data=source.read(item)
                if item.filename=='docProps/core.xml':
                    x=ET.fromstring(data)
                    for name in ('{http://purl.org/dc/elements/1.1/}creator','{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastModifiedBy'):
                        e=x.find(name)
                        if e is None:e=ET.SubElement(x,name)
                        e.text='Klotzkette'
                    data=ET.tostring(x,encoding='utf-8',xml_declaration=True)
                if item.filename=='xl/workbook.xml':
                    x=ET.fromstring(data); names=x.find(ns+'definedNames')
                    if names is None:names=ET.SubElement(x,ns+'definedNames')
                    for e in list(names):
                        if e.get('name') in ('_xlnm.Print_Area','_xlnm.Print_Titles'):names.remove(e)
                    for idx,spec in enumerate(sheets):
                        ET.SubElement(names,ns+'definedName',name='_xlnm.Print_Area',localSheetId=str(idx)).text=f"'{spec['name']}'!{spec['range']}"
                        ET.SubElement(names,ns+'definedName',name='_xlnm.Print_Titles',localSheetId=str(idx)).text=f"'{spec['name']}'!$1:$5"
                    calc=x.find(ns+'calcPr')
                    if calc is None:calc=ET.SubElement(x,ns+'calcPr')
                    calc.attrib.update(calcMode='auto',fullCalcOnLoad='1',forceFullCalc='1')
                    data=ET.tostring(x,encoding='utf-8',xml_declaration=True)
                if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                    x=ET.fromstring(data)
                    setup=x.find(ns+'pageSetup')
                    if setup is None:setup=ET.SubElement(x,ns+'pageSetup')
                    setup.attrib.update(paperSize='9',orientation='landscape',fitToWidth='1',fitToHeight='0')
                    margins=x.find(ns+'pageMargins')
                    if margins is None:margins=ET.SubElement(x,ns+'pageMargins')
                    margins.attrib.update(left='0.3',right='0.3',top='0.35',bottom='0.35',header='0.1',footer='0.1')
                    prop=x.find(ns+'sheetPr')
                    if prop is None:prop=ET.Element(ns+'sheetPr');x.insert(0,prop)
                    ps=prop.find(ns+'pageSetUpPr')
                    if ps is None:ps=ET.SubElement(prop,ns+'pageSetUpPr')
                    ps.set('fitToPage','1')
                    if clear:
                        for c in x.findall('.//'+ns+'c'):
                            if c.find(ns+'f') is not None:
                                for v in c.findall(ns+'v'):c.remove(v)
                    data=ET.tostring(x,encoding='utf-8',xml_declaration=True)
                dest.writestr(item,data)
        file.write_bytes(output.getvalue())
    def formulas(file):
        w=load_workbook(file,data_only=False)
        result={(s.title,c.coordinate):c.value.replace("'",'') for s in w for row in s for c in row if c.data_type=='f'}
        w.close();return result
    binary=os.environ.get('SOFFICE','').strip()
    if not binary or not Path(binary).is_file():raise RuntimeError('SOFFICE muss auf die gebündelte Office-Laufzeit zeigen.')
    def native(paths,folder):
        folder.mkdir(parents=True,exist_ok=True)
        cmd=[binary,f'-env:UserInstallation={(folder/"profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(folder),*[str(p) for p in paths]]
        rc,log=run_office(cmd,env=os.environ.copy(),timeout=180)
        (folder/'office.log').write_text(log)
        if rc:raise RuntimeError(log)
        output=[folder/p.name for p in paths]
        if not all(x.is_file() for x in output):raise RuntimeError(log)
        return output
    baselines=[];old=[]
    for report in reports:
        p=CASE/report['file'];decorate(p,report['sheets'],clear=True);baselines.append(p);old.append(formulas(p))
        artifact_log=CASE/(p.name+'.inspect.ndjson')
        if artifact_log.exists():shutil.move(artifact_log,QA/artifact_log.name)
    results=[]
    for report,p,source,expected_formulas in zip(reports,native(baselines,QA/'native-baseline'),baselines,old):
        actual=formulas(p)
        if actual!=expected_formulas:raise AssertionError('Formeln verändert: '+p.name)
        w=load_workbook(p,data_only=True)
        errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e']
        assert not errors,errors
        end=len(json.loads(DATA.read_text())['invoices'])+6
        spec={'120_Kostenentwicklung.xlsx': [('Kostenstand','D13',3482215.41)],
              '121_Rechnungen_und_Zahlungen.xlsx':[('Offene Posten','E34',-89250)],
              '122_Finanzierung_und_Vermietung.xlsx':[('Finanzierung','B18',102000),('Finanzierung','B22',2400),('Tilgung','D6',4800),('Tilgung','F6',1598000)]}[p.name]
        for sheet,cell,value in spec:assert abs(w[sheet][cell].value-value)<.005,(sheet,cell,w[sheet][cell].value)
        w.close();shutil.copyfile(p,source);decorate(source,report['sheets'])
        results.append(dict(file=p.name,formula_count=len(actual),baseline=spec,mutations=[]))
    mutation_paths=[];mutation_specs=[]
    for report in reports:
        for spec in report['mutations']:
            p=Path(spec['path']);decorate(p,report['sheets'],clear=True);mutation_paths.append(p);mutation_specs.append((report['file'],spec))
    for p,(name,spec) in zip(native(mutation_paths,QA/'native-mutations'),mutation_specs):
        w=load_workbook(p,data_only=True);value=w[spec['outputSheet']][spec['outputCell']].value;w.close()
        assert abs(value-spec['expected'])<.005 if isinstance(spec['expected'],(float,int)) else value==spec['expected'],(spec,value)
        next(x for x in results if x['file']==name)['mutations'].append(dict(sheet=spec['sheet'],cell=spec['cell'],value=spec['value'],result=value))
    (QA/'native-recalculation.json').write_text(json.dumps(results,ensure_ascii=False,indent=2)+'\n')
    print('Native Neuberechnung und acht Eingabemutationen erfolgreich.')


def verify_einvoices():
    """Prüft die 28 E-Rechnungen mit lokal bereitgestelltem offiziellem KoSIT-Paket."""
    import hashlib
    import subprocess
    from decimal import Decimal
    from email import policy
    from email.parser import BytesParser
    package=Path(os.environ.get('KOSIT_HOME',str(QA/'kosit')))
    java_path=os.environ.get('KOSIT_JAVA','').strip()
    if not java_path or not Path(java_path).is_file():raise RuntimeError('KOSIT_JAVA muss auf eine Java-Laufzeit ab Version 11 zeigen.')
    java=Path(java_path)
    config=package/'configuration'
    reports=package/'reports';reports.mkdir(parents=True,exist_ok=True)
    paths=sorted(CASE.glob('*_XRechnung.xml'))
    dataset=json.loads(DATA.read_text())
    ns={'cac':'urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2',
        'cbc':'urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2',
        'rep':'http://www.xoev.de/de/validator/varl/1'}
    def run(paths,output):
        output.mkdir(parents=True,exist_ok=True)
        cmd=[str(java),'-jar',str(package/'validator-1.6.3-standalone.jar'),'-s',str(config/'scenarios.xml'),'-r',str(config),'-o',str(output),'-h',*[str(p) for p in paths]]
        result=subprocess.run(cmd,capture_output=True,text=True,timeout=180)
        (output/'run.log').write_text(result.stdout+result.stderr)
        return result.returncode
    assert len(paths)==28,len(paths)
    assert run(paths,reports)==0,'KoSIT hat eine Originalrechnung abgewiesen.'
    checks=[]
    for p in paths:
        r=ET.parse(p).getroot()
        invoice=next(i for i in dataset['invoices'] if i['number']==int(p.name.split('_')[0]))
        assert r.findtext('cbc:ID',namespaces=ns)==invoice['id']
        assert r.findtext('cbc:IssueDate',namespaces=ns)==invoice['date']
        payable=Decimal(r.findtext('cac:LegalMonetaryTotal/cbc:PayableAmount',namespaces=ns))
        assert payable==abs(Decimal(str(invoice['gross'])))
        credit=r.tag.endswith('CreditNote')
        assert credit==(invoice['gross']<0)
        report=ET.parse(reports/(p.stem+'-report.xml')).getroot()
        assert report.get('valid')=='true' and report.find('rep:assessment/rep:accept',ns) is not None
        steps=[{'id':step.get('id'),'valid':step.get('valid')} for step in report.findall('.//rep:validationStepResult',ns)]
        assert all(step['valid']=='true' for step in steps)
        checks.append(dict(file=p.name,id=invoice['id'],payable=float(payable),credit=credit,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),steps=steps))
    negative=package/'negative';negative.mkdir(parents=True,exist_ok=True)
    probe=ET.parse(paths[0]);amount=probe.getroot().find('cac:LegalMonetaryTotal/cbc:PayableAmount',ns)
    amount.text=str(Decimal(amount.text)+1)
    bad=negative/'wrong-payable.xml';probe.write(bad,encoding='utf-8',xml_declaration=True)
    assert run([bad],negative)!=0,'Fehlerhafte Zahlsumme wurde nicht erkannt.'
    report=ET.parse(negative/'wrong-payable-report.xml').getroot()
    assert report.get('valid')=='false'
    assert 'BR-CO-16' in (negative/'wrong-payable-report.xml').read_text()
    output={'validator':'KoSIT Validator 1.6.3','configuration':'XRechnung 3.0.2 / 2026-08-31',
            'validated_at':'2026-09-27','accepted':len(checks),'rejected':0,'checks':checks,
            'negative_probe':{'mutation':'PayableAmount + 1 EUR','rejected':True,'rule':'BR-CO-16'}}
    (QA/'xrechnung-qa.json').write_text(json.dumps(output,ensure_ascii=False,indent=2)+'\n')
    messages=[]
    for record in dataset['records']:
        if not record['filename'].endswith('.eml'):continue
        file=CASE/record['filename'];message=BytesParser(policy=policy.default).parsebytes(file.read_bytes())
        for field in ('From','To'):
            assert len(message[field].addresses)==1,(file.name,field)
            assert message[field].addresses[0].addr_spec.endswith('.example')
        assert not message.defects,(file.name,message.defects)
        attached={part.get_filename():part.get_payload(decode=True) for part in message.iter_attachments()}
        assert sorted(attached)==sorted(record['attachments'])
        assert all(data==(CASE/name).read_bytes() for name,data in attached.items())
        messages.append(dict(file=file.name,headers='valid',attachments=list(attached)))
    assert len(messages)==4
    (QA/'eml-qa.json').write_text(json.dumps(messages,ensure_ascii=False,indent=2)+'\n')
    print('28 XRechnungen akzeptiert; negative Betragsprobe abgewiesen; vier EML-Dateien vollständig geprüft.')


if __name__=='__main__':
    if '--qa' in sys.argv:verify_workbooks()
    elif '--invoice-qa' in sys.argv:verify_einvoices()
    else:build()
