#!/usr/bin/env python3
"""Regressionsprüfungen für die Zugewinnakte und ihr eigenständiges Plugin."""

import csv
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from testakte_zip_common import working_dump_flat_pairs

ROOT = Path(__file__).resolve().parent.parent
SPEC = importlib.util.spec_from_file_location("bergmann", ROOT / "scripts/build-zugewinn-bergmann-belege.py")
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)
SLUG = BUILDER.SLUG
D = Decimal


def rows(path):
    with path.open(encoding="utf-8") as stream:
        return list(csv.DictReader(stream, delimiter=";"))


class ZugewinnTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.generated = Path(cls.temp.name)
        cls.data = BUILDER.build(cls.generated)
        cls.saved = ROOT / "testakten" / SLUG

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_fresh_and_saved_beleg_inventory(self):
        generated = self.generated / "testakten" / SLUG
        expected = list(generated.rglob("*.pdf"))
        self.assertEqual(len(expected), 77)
        self.assertEqual(len(list(generated.rglob("*.csv"))), 8)
        for source in expected:
            relative = source.relative_to(generated)
            with self.subTest(file=relative):
                saved = self.saved / relative
                self.assertTrue(saved.is_file(), saved)
                for path in (source, saved):
                    document = PdfReader(path)
                    self.assertEqual(len(document.pages), 1, path)
                    text = document.pages[0].extract_text()
                    self.assertGreater(len(text), 600, path)
                    for forbidden in ("Musterlösung", "Lösungsmatrix", "Platzhalter", "Testakte", chr(167)):
                        self.assertNotIn(forbidden, text, path)

    def test_rebuild_is_deterministic(self):
        before = {p.relative_to(self.generated): hashlib.sha256(p.read_bytes()).digest() for p in self.generated.rglob("*") if p.is_file()}
        BUILDER.build(self.generated)
        after = {p.relative_to(self.generated): hashlib.sha256(p.read_bytes()).digest() for p in self.generated.rglob("*") if p.is_file()}
        self.assertEqual(before, after)

    def test_csv_rows_links_and_saved_values(self):
        generated = self.generated / "testakten" / SLUG
        for source in sorted((generated / "zahlenwerk").glob("*.csv")):
            saved = self.saved / "zahlenwerk" / source.name
            self.assertEqual(rows(source), rows(saved), source.name)
            for path in (source, saved):
                with path.open(encoding="utf-8") as stream:
                    contents = list(csv.reader(stream, delimiter=";"))
                self.assertGreaterEqual(len(contents), 5, path.name)
                self.assertTrue(all(len(row) == len(contents[0]) for row in contents), path.name)
                for row in rows(path):
                    self.assertTrue((path.parent.parent / row["Belegdatei"]).is_file(), row)

    def test_bank_running_balances_dates_and_count(self):
        total = 0
        for name, meta in BUILDER.ACCOUNTS.items():
            ledger = rows(self.saved / "zahlenwerk" / f"Konto_{name}_2025_2026.csv")
            self.assertEqual(ledger[0]["Referenz"], "VORTRAG")
            self.assertEqual(D(ledger[0]["Saldo_EUR"]), D(meta["opening"]))
            balance = D(meta["opening"])
            refs = set()
            for row in ledger[1:]:
                balance += D(row["Betrag_EUR"])
                self.assertEqual(balance, D(row["Saldo_EUR"]), row)
                self.assertNotIn(row["Referenz"], refs)
                refs.add(row["Referenz"])
                day = date.fromisoformat(row["Buchungstag"])
                self.assertLessEqual(day, date(2026, 7, 8))
                self.assertGreater(day, date(2025, 4, 1))
                self.assertLess(day.weekday(), 5)
            self.assertEqual(balance, D(self.data["accounts"][name]["closing"]))
            total += len(ledger) - 1
        self.assertEqual(total, 481)

    def test_own_transfers_have_two_opposite_bookings(self):
        paired = defaultdict(list)
        for name in BUILDER.ACCOUNTS:
            for row in rows(self.saved / "zahlenwerk" / f"Konto_{name}_2025_2026.csv"):
                ref = row["Referenz"]
                if ref.endswith(("-M-H", "-J-H", "-J-M")) or ref == "SONDER-202605":
                    paired[ref].append(row)
        self.assertEqual(len(paired), 47)
        for ref, pair in paired.items():
            self.assertEqual(len(pair), 2, ref)
            self.assertEqual(sum(D(row["Betrag_EUR"]) for row in pair), 0, ref)
            self.assertEqual(pair[0]["Buchungstag"], pair[1]["Buchungstag"], ref)

    def test_mortgage_is_not_the_registered_charge(self):
        ledger = rows(self.saved / "zahlenwerk/Darlehensbewegungen_2025_2026.csv")
        balance = D("207860.00")
        self.assertEqual(len(ledger), 16)
        for row in ledger:
            self.assertEqual(D(row["Anfang_EUR"]), balance)
            balance += D(row["Zinsen_EUR"]) - D(row["Rate_EUR"]) - D(row["Sondertilgung_EUR"])
            self.assertEqual(balance, D(row["Restkapital_EUR"]))
        self.assertEqual(balance, D("188007.65"))
        self.assertNotEqual(balance, D("280000"))
        household = rows(self.saved / "zahlenwerk/Konto_Haushalt_2025_2026.csv")
        regular = [r for r in household if r["Referenz"].endswith("-HD")]
        self.assertEqual(len(regular), len(ledger))
        self.assertEqual(-sum(D(r["Betrag_EUR"]) for r in regular), sum(D(r["Rate_EUR"]) for r in ledger))
        self.assertEqual(sum(D(r["Sondertilgung_EUR"]) for r in ledger), D("3500"))

    def test_crypto_partial_sale_and_cash_are_distinct(self):
        ledger = rows(self.saved / "zahlenwerk/Krypto_Transaktionen.csv")
        quantity = D(0)
        for row in ledger:
            quantity += D(row["BTC_Veränderung"])
            self.assertEqual(quantity, D(row["BTC_Bestand"]))
        self.assertEqual(quantity, D("0.04"))
        sale = next(r for r in ledger if r["Vorgang"] == "Verkauf")
        paid = D(sale["EUR_Betrag"]) - D(sale["Gebühr_EUR"])
        bank = rows(self.saved / "zahlenwerk/Konto_Jonas_2025_2026.csv")
        incoming = [r for r in bank if r["Referenz"] == sale["Referenz"]]
        self.assertEqual(len(incoming), 1)
        self.assertEqual(D(incoming[0]["Betrag_EUR"]), paid)
        self.assertEqual(paid, D("1874.36"))

    def test_business_rows_reconcile_without_a_company_value(self):
        for row in rows(self.saved / "zahlenwerk/Offene_Posten_2026-07-08.csv"):
            self.assertEqual(D(row["Brutto_EUR"]) - D(row["Bezahlt_EUR"]), D(row["Rest_EUR"]))
        inventory = rows(self.saved / "zahlenwerk/Inventar_Jonas_2026-07-08.csv")
        self.assertEqual(len(inventory), 8)
        self.assertTrue(any(D(r["Buchwert_EUR"]) != D(r["Angebot_Haendler_EUR"]) for r in inventory))
        self.assertTrue(all("Unternehmenswert" not in r for r in inventory))

    def test_savings_growth_is_not_double_counted_in_cash(self):
        mara = rows(self.saved / "zahlenwerk/Konto_Mara_2025_2026.csv")
        savings = [r for r in mara if r["Referenz"].endswith("-TG")]
        fund = [r for r in mara if r["Referenz"].endswith("-DEP")]
        self.assertEqual(len(savings), 16)
        self.assertEqual(len(fund), 15)
        net_interest = sum(map(D, ("84.60", "88.30", "90.80", "47.32", "51.46")))
        self.assertEqual(D(34600) - sum(D(r["Betrag_EUR"]) for r in savings) + net_interest, D("44562.48"))

    def test_native_correspondence_and_workbooks(self):
        letters = list((self.saved / "aktenstuecke").glob("*.docx"))
        self.assertGreaterEqual(len(letters), 12)
        for path in letters:
            doc = Document(path)
            for section in doc.sections:
                self.assertAlmostEqual(section.page_width.cm, 21, delta=0.1, msg=path.name)
                self.assertAlmostEqual(section.page_height.cm, 29.7, delta=0.1, msg=path.name)
            text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertGreater(len(text), 600, path.name)
            self.assertNotIn(chr(167), text, path.name)
        messages = list((self.saved / "korrespondenz").glob("*.eml"))
        self.assertGreaterEqual(len(messages), 10)
        for path in messages:
            msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            self.assertFalse(msg.defects, path.name)
            for key in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version"):
                self.assertTrue(msg[key], (path.name, key))
            self.assertGreater(len(msg.get_body(preferencelist=("plain",)).get_content()), 200, path.name)
        for filename, tabs in [("Kontoumsaetze_2025_2026.xlsx", ["Mara", "Jonas", "Haushalt"]),
                               ("Betriebsunterlagen_2023_2026.xlsx", ["Jahreszahlen", "OffenePosten", "Inventar"])]:
            wb = load_workbook(self.saved / "xlsx" / filename, data_only=True, read_only=True)
            try:
                self.assertEqual(wb.sheetnames, tabs)
                for sheet in wb:
                    for row in sheet:
                        self.assertTrue(all(c.data_type != "e" for c in row), sheet.title)
            finally:
                wb.close()

    def test_inheritance_dates_and_discharge_match(self):
        bank = " ".join(PdfReader(self.saved / "belege/Erbschaft_Horst_Bank_2019.pdf").pages[0].extract_text().split())
        self.assertIn("Auskunft 29.03.2019", bank)
        self.assertIn("Am 25. März", bank)
        land = " ".join(PdfReader(self.saved / "belege/Grundstueck_Wittbrietzen_2026.pdf").pages[0].extract_text().split())
        self.assertIn("20. Mai 2019 gelöscht", land)
        register = Document(self.saved / "aktenstuecke/14_grundbuch_wittbrietzen_textabschrift_2026-08-24.docx")
        self.assertIn("20.05.2019", "\n".join(p.text for p in register.paragraphs))

    def test_workbook_records_and_formula_caches_match_sources(self):
        books = {
            "Kontoumsaetze_2025_2026.xlsx": {name: f"Konto_{name}_2025_2026.csv" for name in BUILDER.ACCOUNTS},
            "Betriebsunterlagen_2023_2026.xlsx": {"Jahreszahlen": "Euer_Jonas_2023_2025.csv", "OffenePosten": "Offene_Posten_2026-07-08.csv", "Inventar": "Inventar_Jonas_2026-07-08.csv"},
        }
        for filename, tabs in books.items():
            wb = load_workbook(self.saved / "xlsx" / filename, data_only=True)
            formulas = load_workbook(self.saved / "xlsx" / filename, data_only=False)
            try:
                for tab, source in tabs.items():
                    expected = rows(self.saved / "zahlenwerk" / source)
                    sheet = wb[tab]
                    self.assertEqual(sheet.max_row - 9, len(expected), tab)
                    for index, row in enumerate(expected, 10):
                        for column, expected_value in enumerate(row.values(), 1):
                            actual = sheet.cell(index, column).value
                            if isinstance(actual, datetime):
                                self.assertEqual(actual.date().isoformat(), expected_value)
                            elif isinstance(actual, (int, float)):
                                self.assertEqual(D(str(actual)), D(expected_value))
                            else:
                                self.assertEqual(actual, expected_value)
                    for row in formulas[tab]:
                        for cell in row:
                            if cell.data_type == "f":
                                self.assertIsInstance(sheet[cell.coordinate].value, (int, float), (tab, cell.coordinate))
                    if tab in BUILDER.ACCOUNTS:
                        for cell, positive in (("B6", True), ("E6", False)):
                            total = sum(D(r["Betrag_EUR"]) for r in expected if (D(r["Betrag_EUR"]) > 0) == positive)
                            self.assertEqual(D(str(sheet[cell].value)).quantize(D("0.01")), total)
                        self.assertEqual(D(str(sheet["H6"].value)), D(expected[-1]["Saldo_EUR"]))
            finally:
                wb.close()
                formulas.close()

    def test_two_legible_image_sources(self):
        pictures = list((self.saved / "bilder").glob("*.png"))
        self.assertEqual(len(pictures), 2)
        for path in pictures:
            with Image.open(path) as im:
                self.assertGreaterEqual(im.width, 860)
                self.assertGreaterEqual(im.height, 1800)
                self.assertGreater(len(im.getcolors(im.width * im.height)), 100)

    def test_flat_archive_inventory_does_not_ship_editorial_material(self):
        pairs = working_dump_flat_pairs(self.saved, include_gesamt_pdf=False)
        names = [name for _, name in pairs]
        self.assertEqual(len(names), len(set(n.casefold() for n in names)))
        self.assertTrue(all("/" not in n and "\\" not in n for n in names))
        self.assertFalse(any(Path(n).suffix in (".md", ".yaml", ".yml") for n in names))
        for suffix in (".pdf", ".docx", ".xlsx", ".csv", ".eml", ".txt", ".png"):
            self.assertTrue(any(n.endswith(suffix) for n in names), suffix)

    def test_plugin_has_ten_real_skills_and_separate_prompts(self):
        folder = ROOT / "zugewinnausgleich"
        manifest = json.loads((folder / ".claude-plugin/plugin.json").read_text())
        self.assertEqual(manifest["name"], folder.name)
        skills = list((folder / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 10)
        for suffix in ("schnellstart", "hauptproblem"):
            prompt = folder / f"zugewinnausgleich-{suffix}.md"
            self.assertLessEqual(len(prompt.read_bytes()), 7500)
        self.assertGreater(len((folder / "zugewinnausgleich-werkstatt.md").read_bytes()), 12000)
        for skill in skills:
            text = skill.read_text()
            self.assertGreater(len(text), 1800)
            self.assertNotIn("BEGIN", text)


if __name__ == "__main__":
    unittest.main()
