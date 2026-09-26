#!/usr/bin/env python3
"""Drei eigenständige Gebäudeakten. Autor: Klotzkette. Nur phaseneigene Dateien."""
from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import os
from datetime import datetime
from email import policy
from email.message import EmailMessage
from pathlib import Path
import subprocess
import tempfile
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, PngImagePlugin
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, PageBreak

from akten_build_runtime import serif_font_path, screen_font, node_binary
from readme_decimal_headings import normalize_decimal_headings
from testakte_download_notices import ensure_download_notices

ROOT = Path(__file__).resolve().parents[1]
AUTHOR = "Klotzkette"
CASES = {
    7: {"slug": "bauwirtschaft-hoai-7-vergabe-bibliothek-melle", "title": "Leistungsphase 7: Angebotswertung und Vergabe der Bibliothek Hasebogen in Melle", "ref": "HM-26-F01", "place": "Melle", "date": "2026-09-25"},
    8: {"slug": "bauwirtschaft-hoai-8-bauueberwachung-kita-verden", "title": "Leistungsphase 8: Bauüberwachung, Aufmaß und Abnahmemängel der Kita Mühlenwiese in Verden", "ref": "VM-26-D", "place": "Verden", "date": "2026-08-18"},
    9: {"slug": "bauwirtschaft-hoai-9-objektbetreuung-rathaus-uelzen", "title": "Leistungsphase 9: Spätere Mängel, Vorfristbegehung und Sicherheiten am Rathaus Uelzen", "ref": "UR-21-West", "place": "Uelzen", "date": "2026-09-25"},
}
ACTORS = {
    "stiftung": ("Bibliotheksstiftung Hasebogen", "Lesegasse 41, 49324 Melle", "vorstand@hasebogen.example", "Janna Riek, Vorstand"),
    "mplan": ("Konturhaus Architektur", "Papierhof 27, 49324 Melle", "planung@konturhaus.example", "Ole Venn, Architekt"),
    "klee": ("Kleefeld Fensterbau GmbH", "Werkbogen 39, 49326 Melle", "angebot@kleefeld.example", "Mira Kleefeld, Geschäftsführerin"),
    "breden": ("Breden Glas und Rahmen GmbH", "Industriespur 18, 49152 Bad Essen", "kalkulation@breden.example", "Arnd Breden, Geschäftsführer"),
    "fenn": ("Fennwerk Bauelemente GmbH", "Hallenpfad 24, 32257 Bünde", "vertrieb@fennwerk.example", "Silke Fenn, Geschäftsführerin"),
    "foerder": ("Regionalfonds Lesekultur gGmbH", "Kulturzeile 31, 49074 Osnabrück", "foerderung@lesekultur.example", "Marek Otten, Förderreferat"),
    "kita": ("Mühlenwiese Betreuung gGmbH", "Kinderbogen 32, 27283 Verden", "bau@muehlenwiese.example", "Tessa Lühr, Geschäftsführerin"),
    "vplan": ("Rundbogen Architektur", "Planerzeile 29, 27283 Verden", "objekt@rundbogen.example", "Elisa Wenke, Architektin"),
    "dach": ("Allerfirst Dachbau GmbH", "Werkstieg 36, 27308 Kirchlinteln", "bauleitung@allerfirst.example", "Nils Gerdes, Bauleitung"),
    "technik": ("Prüfbüro Querschnitt", "Ingenieurhof 42, 27283 Verden", "fachplanung@querschnitt.example", "Sven Rölke, Fachplaner"),
    "stadt": ("Stadt Uelzen, Projektstelle Rathaus West", "Verwaltungsbogen 37, 29525 Uelzen", "rathausprojekt@uelzen-projekt.example", "Henrike Falk, Projektleitung"),
    "uplan": ("Büro Raumkante", "Zeichenhof 28, 29525 Uelzen", "objektbetreuung@raumkante.example", "Dora Sell, Architektin"),
    "udach": ("Ilmenaudach Bau GmbH", "Dachweg 46, 29525 Uelzen", "service@ilmenaudach.example", "Björn Eilers, Geschäftsführer"),
    "umetall": ("Heideprofil Metallbau GmbH", "Profilring 33, 29549 Bad Bevensen", "service@heideprofil.example", "Maren Knoll, Geschäftsführerin"),
    "wart": ("Dachpflege Suderland GmbH", "Betriebsweg 44, 29556 Suderburg", "disposition@suderland.example", "Leon Tamm, Servicetechniker"),
}
FILES: dict[int, dict[int, dict]] = {7: {}, 8: {}, 9: {}}
FONT = "AktenSerif"
BODY = ParagraphStyle("Text", fontName=FONT, fontSize=11, leading=14, spaceAfter=8)
SMALL = ParagraphStyle("Tabelle", parent=BODY, fontSize=9.5, leading=12, spaceAfter=4)
HEAD = ParagraphStyle("Abschnitt", parent=BODY, fontName=FONT+"Bold", fontSize=12, leading=15, spaceBefore=10, spaceAfter=9, keepWithNext=True)
TITLE = ParagraphStyle("Titel", parent=HEAD, fontSize=16, leading=19, spaceAfter=14)


def euro(value):
    return f"{value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def register_fonts():
    for suffix, bold in [("", False), ("Bold", True)]:
        pdfmetrics.registerFont(TTFont(FONT+suffix, str(serif_font_path(bold))))
    pdfmetrics.registerFontFamily(FONT, normal=FONT, bold=FONT+"Bold", italic=FONT, boldItalic=FONT+"Bold")


def add(phase, number, filename, date, issuer, recipient, title, content, attachments=()):
    FILES[phase][number] = dict(phase=phase, number=number, filename=f"{number:02d}_{filename}", date=date,
                              issuer=issuer, recipient=recipient, title=title, content=content, attachments=list(attachments))


def p(text, style=BODY):
    return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)


def table(rows, widths=None):
    widths = widths or [495/len(rows[0])]*len(rows[0])
    obj = Table([[p(cell, SMALL) for cell in row] for row in rows], colWidths=widths, repeatRows=1)
    obj.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9edef")),
                            ("LINEBELOW", (0, 0), (-1, 0), .6, colors.grey), ("LINEBELOW", (0, 1), (-1, -1), .25, colors.lightgrey),
                            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    return obj


def doc_path(d):
    return ROOT/"testakten"/CASES[d["phase"]]["slug"]/d["filename"]


def attachment_names(d):
    return "; ".join(FILES[d["phase"]][n]["filename"] for n in d["attachments"]) or "keine"


def pdf(d):
    c = CASES[d["phase"]]; sender = ACTORS[d["issuer"]]; recipient = ACTORS[d["recipient"]]
    def frame(canvas, document):
        canvas.setFont(FONT+"Bold", 13); canvas.drawString(50, 800, sender[0])
        canvas.setFont(FONT, 9); canvas.drawString(50, 784, sender[1]); canvas.drawString(50, 771, sender[2])
        canvas.setFont(FONT, 8); canvas.drawString(50, 30, c["ref"]+" | "+d["date"]); canvas.drawRightString(545, 30, f"Seite {document.page}")
    story = [p(recipient[0]+"\n"+recipient[1], SMALL), p(c["place"]+", "+datetime.fromisoformat(d["date"]).strftime("%d.%m.%Y")+" | "+c["ref"], SMALL), p(d["title"], TITLE), p("Sehr geehrte Damen und Herren,")]
    for item in d["content"]:
        if isinstance(item, str): story.append(p(item))
        elif item[0] == "h": story.append(p(item[1], HEAD))
        elif item[0] == "t": story.extend([table(item[1], item[2] if len(item)>2 else None), Spacer(1, 9)])
        elif item[0] == "page": story.append(PageBreak())
    story.append(KeepTogether([Spacer(1, 6), p("Mit freundlichen Grüßen"), p("gez. "+sender[3]), p("Anlagen: "+attachment_names(d), SMALL)]))
    SimpleDocTemplate(str(doc_path(d)), pagesize=A4, leftMargin=50, rightMargin=50, topMargin=94, bottomMargin=49,
                      title=d["title"], author=AUTHOR, creator=AUTHOR).build(story, onFirstPage=frame, onLaterPages=frame)


def docx(d):
    doc = Document(); sec = doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(1.6); sec.bottom_margin=Cm(1.6); sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
    for style in doc.styles:
        if style.type == 1:
            style.font.name="Times New Roman"; style.font.size=Pt(11); style.font.color.rgb=RGBColor(0,0,0)
            style.paragraph_format.space_after=Pt(7)
            fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if key.endswith("Theme"): del fonts.attrib[key]
            for key in ("ascii", "hAnsi", "eastAsia", "cs"): fonts.set(qn("w:"+key), "Times New Roman")
            for border in style.element.xpath("./w:pPr/w:pBdr"):border.getparent().remove(border)
    doc.styles["Title"].font.size=Pt(16); doc.styles["Title"].paragraph_format.space_after=Pt(12)
    doc.styles["Heading 1"].font.size=Pt(12); doc.styles["Heading 1"].font.bold=True
    doc.styles["Heading 1"].paragraph_format.space_before=Pt(9); doc.styles["Heading 1"].paragraph_format.space_after=Pt(9)
    cp=doc.core_properties; cp.author=AUTHOR; cp.last_modified_by=AUTHOR; cp.title=d["title"]; cp.language="de-DE"
    cp.created=datetime.fromisoformat(d["date"]); cp.modified=cp.created
    sender=ACTORS[d["issuer"]]; recipient=ACTORS[d["recipient"]]; c=CASES[d["phase"]]
    doc.add_paragraph(sender[0]).runs[0].bold=True
    doc.add_paragraph(sender[1]+"\n"+sender[2]); doc.add_paragraph(recipient[0]+"\n"+recipient[1])
    doc.add_paragraph(c["place"]+", "+datetime.fromisoformat(d["date"]).strftime("%d.%m.%Y")+" | "+c["ref"])
    doc.add_paragraph(d["title"], "Title"); doc.add_paragraph("Sehr geehrte Damen und Herren,")
    for item in d["content"]:
        if isinstance(item, str): doc.add_paragraph(item).paragraph_format.keep_together=True
        elif item[0] == "h": doc.add_paragraph(item[1], "Heading 1")
        elif item[0] == "page": doc.add_page_break()
        elif item[0] == "t":
            t=doc.add_table(rows=0, cols=len(item[1][0])); t.style="Table Grid"
            for i, row in enumerate(item[1]):
                cells=t.add_row().cells; t.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                if i == 0: t.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                for cell, value in zip(cells, row):
                    cell.text=str(value)
                    for para in cell.paragraphs:
                        for run in para.runs: run.font.size=Pt(10); run.bold=i == 0
    end=doc.add_paragraph("Mit freundlichen Grüßen\ngez. "+sender[3]); end.paragraph_format.keep_together=True
    doc.add_paragraph("Anlagen: "+attachment_names(d))
    footer=sec.footer.paragraphs[0]; footer.text=c["ref"]+" | "+d["date"]+" | Seite "
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    doc.save(doc_path(d))


def eml(d):
    sender=ACTORS[d["issuer"]]; recipient=ACTORS[d["recipient"]]
    msg=EmailMessage(policy=policy.SMTP); msg["From"]=f"{sender[3]} <{sender[2]}>"; msg["To"]=f"{recipient[0]} <{recipient[2]}>"
    msg["Date"]=datetime.fromisoformat(d["date"]+"T10:35:00+02:00")
    msg["Message-ID"]=f"<hoai{d['phase']}.{d['number']}.{d['date']}@akten.example>"; msg["Subject"]=CASES[d["phase"]]["ref"]+" | "+d["title"]
    body="Sehr geehrte Damen und Herren,\n\n"+"\n\n".join(d["content"])+"\n\nMit freundlichen Grüßen\n"+sender[3]+"\n"+sender[0]+"\n"+sender[1]+"\n\nAnlagen: "+attachment_names(d)+"\n"
    msg.set_content(body)
    for n in d["attachments"]:
        file=doc_path(FILES[d["phase"]][n]); ext=file.suffix
        main, sub=("image", "png") if ext==".png" else ("application", "pdf")
        msg.add_attachment(file.read_bytes(), maintype=main, subtype=sub, filename=file.name)
    doc_path(d).write_bytes(bytes(msg))


Q7=[1,12,8,2,20,72,6,1]
EP7=[1800,950,1350,4200,95,28,420,900]
OFFERS7={"klee":[1950,980,1420,4350,98,30,440,950], "breden":[1650,1010,1380,4200,100,29,410,900], "fenn":[2000,990,1450,4300,96,31,415,1000]}
LV7=[
    ("1.01", "Baustelleneinrichtung", "psch", "Einrichten und Räumen für den Fensteraustausch in zwei Bauabschnitten. Staubwände zur geöffneten Kinderbibliothek, Schutz vorhandener Regale, Hubgerät und tägliche Reinigung sind enthalten. Lagerfläche im Hof maximal 25 m². Vorhaltung für zehn Arbeitstage ist enthalten."),
    ("1.02", "Fenster W1", "St", "Zwölf Holz-Aluminiumfenster, Rohbauöffnung jeweils 1,20 m breit und 1,50 m hoch, ein Drehkippflügel. Dreifachverglasung, geforderter Uw-Wert höchstens 0,95 W/(m²K) für das gesamte Element. Innen Eiche lasiert, außen RAL 7039. Ausbau, Entsorgung, Befestigung und dreiseitiger innerer und äußerer Anschluss sind im Preis enthalten."),
    ("1.03", "Fenster W2", "St", "Acht zweiteilige Holz-Aluminiumfenster, Rohbauöffnung 1,80 m breit und 1,50 m hoch, ein Festfeld und ein Drehkippflügel. Gleiche Oberflächen und gleicher Uw-Höchstwert wie W1. Keine Änderung der Bestandsstürze. Die Anschlussfuge wird nach F-04 Revision B hergestellt."),
    ("1.04", "Eingangselement T1", "St", "Zwei thermisch getrennte Aluminiumtürelemente, Rohbauöffnung 1,30 m breit und 2,40 m hoch. Lichte Durchgangsbreite mindestens 1,00 m. Verglasung im Türbereich als Verbundsicherheitsglas gemäß vereinbarter Produktbeschreibung. Schwellenhöhe 20 mm. Türschließer und Zylinderaufnahme sind enthalten; Schließzylinder liefert der Auftraggeber."),
    ("1.05", "Außenfensterbänke", "St", "Zwanzig beschichtete Aluminiumfensterbänke mit Endstücken und Entdröhnung, Ausladung 250 mm. Zwölf Stück Länge 1,30 m, acht Stück Länge 1,90 m. Lieferung und Einbau einschließlich der im Plan gezeichneten seitlichen Anschlüsse."),
    ("1.06", "Laibungsputz", "m", "Wiederherstellen der inneren Putzanschlüsse entlang der Fenster- und Türlaibungen, mittlere Breite 120 mm. Abrechnung nach laufendem Meter fertigem Anschluss. Grundierung und weißer Schlussanstrich sind enthalten; Vollflächenanstrich der Wände ist nicht enthalten."),
    ("1.07", "Sonnenschutz R1", "St", "Sechs außenliegende Raffstores für die südlichen W2-Elemente, elektrisch, mit Endlageneinstellung. Mechanische Montage und Leitung bis Anschlussdose sind enthalten. Netzanschluss und Gebäudeautomation werden vom gesonderten Elektroauftragnehmer hergestellt. Gemeinsamer Funktionstermin ist einzukalkulieren."),
    ("1.08", "Funktionsprüfung und Einweisung", "psch", "Prüfung sämtlicher Flügel, Beschläge, Türschließer und Raffstores im gemeinsamen Termin. Einweisung des Bibliothekspersonals, Lieferung der Produkt- und Pflegeunterlagen sowie Zuordnung zu den Elementkennungen. Ein zweiter Termin wegen eines vom Auftragnehmer zu vertretenden Mangels ist enthalten."),
]


def define7():
    a=lambda n,f,dt,i,r,t,c,att=(): add(7,n,f,dt,i,r,t,c,att)
    a(1,"Planervertrag.docx","2026-05-18","stiftung","mplan","Architektenvertrag, Umbau Bibliothek Hasebogen",[
        "Die Bibliotheksstiftung Hasebogen und Konturhaus Architektur schließen für den Umbau des Leseflügels an der Lesegasse 41 in Melle folgenden Vertrag. Der Altbau bleibt im Eigentum der Stiftung. Gegenstand sind Gebäudeplanung und Koordination des Austauschs der Fenster und Eingangselemente; die Fachplanung Elektrotechnik wird gesondert beauftragt.",
        ("h","1. Vertragsleistung und Abruf"),
        "Die Leistungsphasen 5 bis 7 für Gebäude werden für das Fenster- und Türlos verbindlich abgerufen. Die vorhandene genehmigte Gestaltung wird beibehalten. Konturhaus stellt die Ausführungsdetails und Vergabeunterlagen her, prüft die eingegangenen Angebote einschließlich positionsbezogenem Preisspiegel, führt zulässige Aufklärungsgespräche und erstellt den Vergabevorschlag mit vollständigen Vertragsanlagen. Eine Beauftragung der Leistungsphasen 8 und 9 bedarf eines weiteren Abrufs.",
        ("h","2. Honorar und Zuarbeit"),
        "Für die genannten Leistungen wird ein Pauschalhonorar von 9.600,00 EUR netto zuzüglich 19 Prozent Umsatzsteuer vereinbart. 3.000,00 EUR netto entfallen auf Ausführungsdetails, 3.600,00 EUR netto auf die Vergabevorbereitung und 3.000,00 EUR netto auf die Mitwirkung bei der Vergabe. Erforderliche zusätzliche Planungsvarianten und die Vertretung in Nachprüfungsverfahren sind nicht enthalten. Die Stiftung liefert Satzung, Finanzierungsunterlagen und etwaige Förderbedingungen bis zum Beginn der Angebotsabfrage.",
        ("h","3. Erklärungen und Termine"),
        "Konturhaus darf technische Auskünfte im abgestimmten Umfang vorbereiten. Aufträge und Preisänderungen werden ausschließlich von Janna Riek und dem zweiten Vorstandsmitglied Leonhard Behr gemeinsam erklärt. Eine Zuschlags-, Zahlungs- oder Vergleichsvollmacht wird nicht erteilt. Der Fenstertausch soll vom 16. bis 27. November 2026 erfolgen. Die Bibliothek ist in dieser Zeit abschnittsweise geöffnet. Beide Parteien erhalten eine gleichlautende Ausfertigung.",
        "Für die Stiftung: gez. Janna Riek und Leonhard Behr. Für Konturhaus Architektur: gez. Ole Venn. Melle, 18.05.2026."
    ])
    a(2,"Stiftungsbeschluss.pdf","2026-06-02","stiftung","mplan","Beschlussausfertigung 06/2026, Finanzierung Leseflügel",[
        "Der Vorstand hat in seiner Sitzung am 2. Juni 2026 mit beiden Stimmen die Sanierung des Leseflügels beschlossen. Das beschlossene Gesamtbudget beträgt 1.840.000,00 EUR netto für Gebäude, Technik, Ausstattung und Planung. Für das Fenster- und Türlos stehen 45.000,00 EUR netto zur Verfügung. Der Beschluss ersetzt keinen Einzelauftrag.",
        ("h","1. Finanzierung"),
        "1.590.000,00 EUR sollen aus dem Stiftungsvermögen und bereits zugesagten privaten Spenden aufgebracht werden. Ein Zuschuss über 250.000,00 EUR wurde beim Regionalfonds Lesekultur beantragt; ein Bewilligungsschreiben liegt noch nicht vor. Das Projekt wird nicht durch kommunale Kreditgarantien abgesichert. Die Finanzierung bleibt nach Auskunft der Kassenführung auch bei Ablehnung des Zuschusses gesichert.",
        ("h","2. Organisation"),
        "Die Stiftung ist eine rechtsfähige Stiftung bürgerlichen Rechts. Von sieben Kuratoriumsmitgliedern werden zwei durch die Stadt und fünf durch die privaten Stifter berufen. Der Vorstand besteht aus Janna Riek und Leonhard Behr; beide wurden vom Kuratorium gewählt. Im Geschäftsjahr 2025 betrugen öffentliche laufende Zuwendungen 18 Prozent der Einnahmen. Eine kommunale Weisungsbefugnis für Einzelvergaben ist in der Satzung nicht vorgesehen.",
        ("h","3. Beschaffungsablauf"),
        "Für das Fensterlos sollen drei schriftliche Angebote zu einheitlichen Unterlagen eingeholt werden. Im Anschreiben werden Preis als Vergleichskriterium, feste Mindestanforderungen und Verzicht auf Preisrunden genannt. Falls der Regionalfonds vor Auftragserteilung besondere Auflagen mitteilt, werden diese dem Vorstand vorgelegt. Die Ausführungsplanung muss den Betrieb der Kinderbibliothek im östlichen Flügel ermöglichen.",
        "Anwesend: Janna Riek, Leonhard Behr und Protokollführerin Fenja Linde. Die Niederschrift wurde am 05.06.2026 genehmigt. gez. Janna Riek; gez. Leonhard Behr; gez. Fenja Linde."
    ])
    a(3,"Angebotsaufforderung.docx","2026-08-24","stiftung","klee","Angebotsaufforderung Fenster- und Türlos HM-26-F01",[
        "Wir bitten um Ihr verbindliches Angebot zum Austausch der im beigefügten Leistungsverzeichnis bezeichneten Elemente. Dieselben Unterlagen gehen an Breden Glas und Rahmen GmbH sowie Fennwerk Bauelemente GmbH. Angebote sind bis 14.09.2026, 12:00 Uhr, an vorstand@hasebogen.example einzureichen. Sie bleiben bis 09.10.2026, 24:00 Uhr, bindend.",
        ("h","1. Angebotsinhalt"),
        "Bieten Sie sämtliche acht Positionen mit Einheitspreisen, Positionssummen, Nachlässen sowie Netto-, Steuer- und Bruttosumme an. Bestätigen Sie die Vertragsunterlagen, die Ausführung in zehn Arbeitstagen und die geforderten Produkteigenschaften. Eigene Vertragsbedingungen und Nebenangebote sind nicht vorgesehen. Bieterfragen sind bis 04.09.2026 zu stellen; Antworten werden allen drei Unternehmen gleichlautend übermittelt.",
        ("h","2. Vergleich und Vertragsbedingungen"),
        "Unter den Angeboten, welche die festgelegten Leistungen, Merkmale und Termine erfüllen, wird nach dem Nettopreis einschließlich unbedingter Nachlässe entschieden. Bedingte Zahlungsnachlässe werden nicht in den Vergleichspreis eingerechnet. Eine nachträgliche Preisrunde ist nicht vorgesehen. Technische Unklarheiten dürfen erläutert werden, ohne ein neues Leistungsangebot zu eröffnen.",
        "Vertragsgrundlage sollen das Auftragsschreiben, diese Bedingungen, die gleichlautenden Bieterantworten, das ausgefüllte Langtext-LV vom 24.08.2026 und Plan F-04 Revision B vom 21.08.2026 werden. Es gelten die gesetzlichen werkvertraglichen Regelungen des BGB; VOB/B wird nicht vereinbart. Abgerechnet wird nach tatsächlich ausgeführten Mengen zu den angebotenen Einheitspreisen. Die Schlussrechnung wird nach vollständiger Leistung und Abnahme innerhalb von 30 Tagen bezahlt. Abschläge bedürfen prüfbarer Leistungsnachweise.",
        ("h","3. Ausführung und Abnahme"),
        "Ausführung ist vom 16. bis 27.11.2026 vorgesehen. Der östliche Bibliotheksbetrieb ist staubgeschützt aufrechtzuerhalten; der westliche Rettungsweg muss täglich frei bleiben. Die gemeinsame Abnahme wird mit der Stiftung vereinbart. Für Mängelansprüche wird eine Frist von fünf Jahren ab Abnahme der jeweiligen Vertragsleistung vereinbart. Eine Vertragserfüllungs- oder Mängelsicherheit wird für dieses Los nicht verlangt."
    ],(4,5))
    texts=["Leistungsstand 24.08.2026. Ausführungsgrundlage ist Plan F-04 Revision B vom 21.08.2026. Die Mengen sind für das vollständige Los angegeben; die Elementbezeichnungen verbinden die Positionen mit dem Plan. Alle Positionen umfassen Lieferung und Montage, soweit der jeweilige Text nicht ausdrücklich etwas anderes bestimmt."]
    for index,(pos,title,unit,desc) in enumerate(LV7):
        texts.extend([("h",pos+". "+title), desc+f" Ausgeschriebene Menge: {Q7[index]} {unit}."])
    texts.append("Die Einheitspreise sind netto in EUR anzugeben. Angaben zur Elementleistung beziehen sich auf das gesamte angebotene Element. Einzelne Glaswerte ersetzen die geforderte Elementangabe nicht. Stand und Bezeichnungen bleiben in den Angeboten unverändert.")
    a(4,"Leistungsverzeichnis.pdf","2026-08-24","mplan","stiftung","Leistungsverzeichnis Fenster und Türen, acht Positionen",texts)
    a(5,"Fensterplan_F04_B.png","2026-08-21","mplan","stiftung","F-04 Revision B, Ansichten und Elementtypen",[])
    a(6,"Planerpreise.xlsx","2026-08-24","mplan","stiftung","Bepreistes Planer-LV HM-26-F01",[])
    for n,bidder,date in [(7,"klee","2026-09-11"),(8,"breden","2026-09-14"),(9,"fenn","2026-09-12")]:
        eps=OFFERS7[bidder]; total=sum(q*ep for q,ep in zip(Q7,eps)); discount=.01 if bidder=="fenn" else 0
        net=round(total*(1-discount),2)
        rows=[["Pos.","Leistung", "Menge", "EP EUR", "GP EUR"]]+[[LV7[i][0],LV7[i][1],f"{Q7[i]} {LV7[i][2]}",euro(ep),euro(Q7[i]*ep)] for i,ep in enumerate(eps)]
        content=[f"Wir bieten die vollständigen Leistungen des Loses HM-26-F01 nach Ihren Unterlagen vom 24.08.2026 und der gleichlautenden Antwort vom 07.09.2026 an. Das Angebot wird durch {ACTORS[bidder][3]} erklärt. Es besteht Bindung bis 09.10.2026, 24:00 Uhr. Sämtliche acht Positionen werden in Eigenleistung ausgeführt; eigene AGB werden nicht beigefügt.",("h","1. Preisblatt"),("t",rows,[42,174,73,98,108]),
                 f"Summe vor Nachlass: {euro(total)} EUR netto. Unbedingter Nachlass: {euro(discount*100)} Prozent. Angebotssumme nach unbedingtem Nachlass: {euro(net)} EUR netto. Umsatzsteuer 19 Prozent: {euro(round(net*.19,2))} EUR. Angebotssumme brutto: {euro(round(net*1.19,2))} EUR.",("h","2. Leistungsumfang")]
        for pos,title,unit,desc in LV7: content.append(pos+" "+title+": "+desc)
        if bidder=="breden":
            content.extend([("h","3. Produkte, Zahlung und Ausführung"),"Für W1 und W2 bieten wir das System BG 78 an. Das beiliegende Produktblatt ist Bestandteil dieser Urkunde: Rahmen Holz-Aluminium, Dreifachglas, Ug 0,60 W/(m²K), Referenzfenster Uw 0,98 W/(m²K) bei 1,23 m mal 1,48 m. Für das Projekt wird im Angebotsformular die Anforderung Uw höchstens 0,95 bestätigt. Die Berechnung der projektspezifischen Elemente erfolgt nach Auftrag.",
            "Bei Zahlung binnen zehn Kalendertagen nach Zugang der prüfbaren Rechnung gewähren wir zwei Prozent Skonto. Unter dieser Zahlungsannahme ergibt sich ein Bruttobetrag von "+euro(round(net*1.19*.98,2))+" EUR. Der vorstehende ungekürzte Bruttobetrag bleibt die Angebotssumme. Wir führen vom 16. bis 27.11.2026 aus. Eine technische Besprechung kann am 22.09.2026 stattfinden."])
        elif bidder=="fenn":
            content.extend([("h","3. Produkte und Ausführungsdauer"),"W1 und W2 werden als FW 82 mit nachgewiesenem Elementwert Uw 0,93 W/(m²K) für beide angebotenen Größen ausgeführt. T1 wird als FT 90 geliefert. Die Oberflächen entsprechen den Vergabeunterlagen. Der unbedingte Nachlass gilt auf alle acht Positionen.",
            "Für die Montage benötigen wir zwölf Arbeitstage. Der Beginn am 16.11.2026 ist zugesagt, Fertigstellung ist daher der 01.12.2026. Die Bibliothek kann die offenen Bereiche abschnittsweise nutzen. Eine Verkürzung auf zehn Arbeitstage ist in dieser Kalkulation nicht berücksichtigt."])
        else:
            content.extend([("h","3. Produkte und Termine"),"W1 und W2 werden im System KF 80 gefertigt. Für W1 wird Uw 0,92 W/(m²K), für W2 Uw 0,94 W/(m²K) zugesagt; die zugrunde liegenden Elementberechnungen sind Bestandteil der Herstellung und werden vor Fertigung vorgelegt. T1 wird als KA 90 mit den geforderten Durchgangsmaßen und Verglasungen geliefert.",
            "Wir bestätigen die Ausführung vom 16. bis 27.11.2026 mit vier Monteuren. In der ersten Woche wird der westliche, in der zweiten Woche der südliche Abschnitt bearbeitet. Das Schutzkonzept aus Position 1.01 ist enthalten. Ein bedingter Zahlungsnachlass wird nicht angeboten."])
        content.extend([("h","4. Unternehmenserklärung"),"Unser Betrieb ist für Fenster- und Türelementmontage eingerichtet und beschäftigt zwölf Monteure sowie zwei technische Sachbearbeiter. Vergleichbare abgeschlossene Projekte sind der Lesesaal Lindenhof 2024 mit 24 Elementen und das Verwaltungsgebäude Südhof 2025 mit 31 Elementen. Die Ansprechpartner werden auf gesonderte Anforderung benannt. Eine Haftpflichtdeckung über 3 Mio. EUR für Personen- und Sachschäden besteht. Die Ausführung wird nicht von einer noch ausstehenden Personalbeschaffung abhängig gemacht."])
        a(n,f"Angebot_{bidder}.pdf",date,bidder,"stiftung","Verbindliches Angebot HM-26-F01",content)
    a(10,"Eingangsniederschrift.docx","2026-09-14","stiftung","mplan","Niederschrift Angebotszugang, 14.09.2026",[
        "Janna Riek und Fenja Linde haben am 14.09.2026 um 12:15 Uhr den Eingang im Projektpostfach gemeinsam festgestellt. Die Anhänge wurden in dem folgenden Stand unverändert abgelegt. Die in der Tabelle genannten Beträge sind vom jeweiligen Deckblatt abgeschrieben; eine rechnerische oder technische Prüfung erfolgte bei dieser Aufnahme nicht.",
        ("t",[["Absender","Eingang MESZ","Datei","Deckblatt brutto EUR"],["Kleefeld","11.09.2026 16:42","07_Angebot_klee.pdf",euro(sum(q*e for q,e in zip(Q7,OFFERS7['klee']))*1.19)],["Breden","14.09.2026 10:18","08_Angebot_breden.pdf",euro(sum(q*e for q,e in zip(Q7,OFFERS7['breden']))*1.19)],["Fennwerk","14.09.2026 08:06","09_Angebot_fenn.pdf",euro(round(sum(q*e for q,e in zip(Q7,OFFERS7['fenn']))*.99,2)*1.19)]],[88,95,183,129]),
        "Es gingen drei PDF-Angebote mit jeweils vollständigem Preisblatt und Unternehmenserklärung ein. Die Breden-Datei enthält einen technischen Produktabschnitt und eine gesonderte Skontoangabe. Fennwerk nennt im letzten Abschnitt eine Montagezeit von zwölf Arbeitstagen. Nach 12:00 Uhr wurde bis zum Abschluss um 12:38 Uhr kein weiteres Angebot festgestellt.",
        "Das Postfach wird ausschließlich durch die Geschäftsstelle verwaltet. Eine Portalöffnung oder elektronische Signaturprüfung fand nicht statt. Die Originalnachrichten verbleiben im Postfacharchiv der Stiftung. Diese Niederschrift bestätigt nur den hier dokumentierten Eingang. gez. Janna Riek; gez. Fenja Linde."
    ])
    a(11,"Bieterfrage_Raffstores.eml","2026-09-03","klee","mplan","Raffstores und Uw-Angabe",[
        "Zum LV vom 24.08.2026 haben wir zwei Fragen. Bei Position 1.07 verstehen wir die sechs Raffstores als die sechs südlichen W2-Elemente. Die beiden weiteren W2 liegen nach Plan auf der Westseite und erhalten keinen Sonnenschutz. Bitte bestätigen Sie diese Zuordnung.",
        "Bei W1 und W2 wird Uw 0,95 gefordert. Wir gehen davon aus, dass dies für die tatsächlich ausgeschriebenen Elementgrößen und nicht nur für ein Referenzfenster gilt. Die Anschlusssituation nach F-04 B ist in unserer Kalkulation vorgesehen. Eine geänderte Verglasung als Nebenangebot beabsichtigen wir nicht."
    ])
    a(12,"Gleichlautende_Antwort.eml","2026-09-07","stiftung","mplan","Weiterleitung der Antwort an alle drei Anbieter",[
        "Die nachstehende Antwort wurde heute um 09:20 Uhr einzeln und gleichlautend an Kleefeld, Breden und Fennwerk versandt. Eine Änderung der Angebotsfrist erfolgt nicht. Bitte verwenden Sie diese Fassung im Vertragsunterlagensatz.",
        "Position 1.07 betrifft die sechs südlichen W2-Elemente; die zwei westlichen W2 erhalten keine Raffstores. Der geforderte Uw-Wert höchstens 0,95 W/(m²K) gilt für die ausgeschriebenen Elementgrößen W1 und W2 als Ganzes. Ein Ug-Wert des Glases oder ein Wert für eine andere Referenzgröße ersetzt diese Angabe nicht. Mengen und übrige Bedingungen bleiben unverändert.",
        "Die Absendebelege mit den drei Empfängern sind im Postfachordner HM-26-F01 abgelegt. Alle drei Firmen haben den Empfang am selben Tag per Kurzantwort bestätigt."
    ])
    a(13,"Breden_Erlaeuterung.eml","2026-09-22","breden","mplan","BG 78 und geplanter Fertigungsnachweis",[
        "Zu Ihrer telefonischen Frage vom 21.09.2026 bestätigen wir, dass unser Angebot vom 14.09.2026 unverändert das System BG 78 enthält. Der Wert 0,98 stammt aus einem älteren Referenzblatt. Die projektspezifische Berechnung liegt uns noch nicht vor; wir erwarten sie am 29.09.2026.",
        "Unser Lieferant hat mündlich erklärt, mit einer anderen warmen Kante könne voraussichtlich 0,94 erreicht werden. Ob diese Ausführung schon in unserer Kalkulation enthalten war, klärt der Einkauf derzeit. Wir bitten, zunächst weiterhin mit dem eingereichten Preis zu rechnen. Ein neues Preisblatt senden wir heute nicht.",
        "Der Zahlungsnachlass bleibt an zehn Kalendertage gebunden. Unsere Montagekapazität für den angebotenen Zeitraum ist reserviert. Diese Nachricht enthält keine Anlage."
    ])
    a(14,"Foerderstand.eml","2026-09-23","foerder","stiftung","Antrag LF-26041, noch kein Bewilligungsbescheid",[
        "Ihr Antrag über 250.000,00 EUR befindet sich weiter in der Prüfung. Das Fachgremium tagt am 06.10.2026. Diese Nachricht stellt weder eine Bewilligung noch die Zustimmung zu einem vorzeitigen Maßnahmenbeginn dar.",
        "Für geförderte Bauleistungen werden dem Bescheid Beschaffungsbedingungen beigefügt. Ob für Ihr Vorhaben eine vorgezogene Angebotsabfrage zulässig war und welche Dokumentationsanforderungen gelten, wird im laufenden Verfahren geprüft. Bitte übersenden Sie vor einer Bindung gegenüber einem Unternehmer den bisherigen Ablauf an unser Förderreferat.",
        "Der Fonds ist eine gemeinnützige Gesellschaft mit kommunalen und privaten Gesellschaftern. Für die Entscheidung über Ihren Antrag gelten die bei Antragstellung übermittelten Fördergrundsätze. Eine telefonische Auskunft aus Juni ersetzt die schriftliche Entscheidung nicht."
    ])
    a(15,"Termine_Schnittstellen.csv","2026-09-24","mplan","stiftung","Terminabstimmung Konturhaus",[
        ["Vorgang","Beginn","Ende","Beteiligter","Bezug","Bemerkung"],
        ["Vorstandssitzung","2026-09-30","2026-09-30","Stiftung","Beschluss Juni","Förderstand noch offen"],
        ["Bindefrist Angebote","2026-09-14","2026-10-09","drei Anbieter","Angebotsaufforderung","24:00 Uhr"],
        ["Fenstertausch West","2026-11-16","2026-11-20","noch offen","F-04 B","Westlicher Fluchtweg täglich frei"],
        ["Fenstertausch Süd","2026-11-23","2026-11-27","noch offen","F-04 B","Kinderbibliothek Ost geöffnet"],
        ["Elektroanschluss Raffstores","2026-11-26","2026-11-27","Elektro Wiek","separater Vertrag E-26-04","Dosen vorhanden; 6 Motorleitungen"],
        ["Einweisung","2026-11-27","2026-11-27","Bibliotheksleitung","LV 1.08","Beginn 14:00 Uhr vorgesehen"],
    ])
    a(16,"Vorstand_Kostenrahmen.eml","2026-09-25","stiftung","mplan","Vorstandssitzung und Preisangaben",[
        "Herr Behr hat für die Sitzung am 30.09.2026 die Breden-Summe nach Skonto notiert. Unsere Kasse kann Zahlung binnen zehn Tagen jedoch nicht verbindlich zusagen, weil Herr Behr im November zwei Wochen abwesend ist. Über 45.000 EUR netto für das Los soll ohne neuen Beschluss nicht verfügt werden.",
        "Die Bibliotheksleitung hält am Wiedereröffnungstermin 30.11.2026 fest; eine spätere Fertigstellung würde die bereits gebuchte Eröffnungswoche treffen. Über eine Verschiebung wurde noch nicht entschieden. Der Fonds hat gestern erneut darauf hingewiesen, dass vor Auftragserteilung seine schriftliche Rückmeldung abzuwarten sei.",
        "Es wurde bislang kein Auftrag erklärt. Herr Behr hatte angeregt, die günstigste Firma telefonisch um weitere fünf Prozent zu bitten. Ich habe ihm mitgeteilt, dass unsere Aufforderung keine Preisrunde vorsieht. Die nächste gemeinsame Vorstandsunterschrift wäre am 30.09.2026 möglich."
    ])


Q8=[1,216,216,216,64,8,2,1]
EP8=[3200,18,72,39,48,165,440,650]
LV8=[("1.01","Einrichtung und Schutz","psch"),("1.02","Dampfsperre","m²"),("1.03","Gefälledämmung","m²"),("1.04","Abdichtung zweilagig","m²"),("1.05","Randanschlüsse","m"),("1.06","Dachabläufe","St"),("1.07","Notüberläufe","St"),("1.08","Prüfung und Dokumentation","psch")]


def define8():
    a=lambda n,f,dt,i,r,t,c,att=(): add(8,n,f,dt,i,r,t,c,att)
    a(1,"Objektueberwachungsvertrag.docx","2026-03-02","kita","vplan","Abruf Objektüberwachung, Kita Mühlenwiese",[
        "Mühlenwiese Betreuung gGmbH ruft für den eingeschossigen Erweiterungsbau mit zwei Gruppenräumen am Kinderbogen 32 in Verden die Gebäude-LPH 8 aus dem Architektenvertrag vom 12.08.2025 ab. Rundbogen Architektur übernimmt die vereinbarten Überwachungs- und Dokumentationsaufgaben einschließlich gemeinsamen Aufmaßes, Rechnungsprüfung, Kostenfeststellung, Abnahmeorganisation, Übergabe und Verfolgung der bei Abnahme festgestellten Mängel.",
        ("h","1. Zuständigkeiten"),
        "Elisa Wenke ist Ansprechpartnerin der Objektüberwachung. Die öffentlich-rechtliche Bauleiterbestellung liegt bei Bauingenieur Jan Toben auf Grundlage gesonderter Bestellung der Bauherrin. Die Prüfung des Tragwerks erfolgt durch das gesondert beauftragte Tragwerksbüro; die Fachüberwachung der technischen Anlagen durch Prüfbüro Querschnitt. Rundbogen koordiniert die Beiträge, übernimmt aber keine fremden Prüfbescheinigungen als eigene.",
        ("h","2. Kontrolle und Dokumentation"),
        "Kontrolltermine sind nach Bauablauf und erforderlichen Prüfungen vor Verdeckung abzustimmen. Eine tägliche dauernde Anwesenheit wird nicht vereinbart. Die Berichte müssen erkennen lassen, wer tatsächlich vor Ort war und welche Bereiche untersucht wurden. Die Bauherrin erhält wöchentlich den Termin- und Kostenstand sowie unverzüglich Mitteilungen über erhebliche Abweichungen. Ein gesonderter Zahlungsplan und betriebliche Buchhaltung sind nicht beauftragt.",
        ("h","3. Honorar und Vollmacht"),
        "Das Pauschalhonorar beträgt 28.000,00 EUR netto zuzüglich 19 Prozent Umsatzsteuer. Preisänderungen, zusätzliche Unternehmeraufträge, Abnahmeerklärungen und Zahlungen bleiben Tessa Lühr vorbehalten. Rundbogen darf Prüf- und Abstimmungsschreiben vorbereiten und nach gesonderter Freigabe versenden. LPH 9 ist nicht abgerufen. Die Geschäftsführerin zeichnet am 02.03.2026; Elisa Wenke bestätigt den Abruf am selben Tag."
    ])
    rows=[["Pos.","Leistung","Menge","EP EUR","GP EUR"]]+[[pos,t,f"{q} {unit}",euro(ep),euro(q*ep)] for (pos,t,unit),q,ep in zip(LV8,Q8,EP8)]
    a(2,"Dachbauvertrag.docx","2026-04-10","kita","dach","Bauvertrag Dach Erweiterung, VM-26-D",[
        "Mühlenwiese Betreuung gGmbH beauftragt Allerfirst Dachbau GmbH mit dem vollständigen Dachaufbau des Erweiterungsbaus. Die Grundfläche beträgt 18,00 m mal 12,00 m. Es gelten die folgenden acht Positionen sowie Plan D-21 Revision B vom 02.04.2026. VOB/B wird nicht einbezogen. Die Abrechnung erfolgt nach tatsächlichen geometrischen Mengen; Randanschlüsse werden zusätzlich nach Länge abgerechnet, aber nicht nochmals zur ebenen Dachfläche addiert.",
        ("h","1. Preisvereinbarung"),("t",rows,[40,180,65,95,115]),
        f"Auftragssumme netto {euro(sum(q*ep for q,ep in zip(Q8,EP8)))} EUR, Umsatzsteuer 19 Prozent {euro(round(sum(q*ep for q,ep in zip(Q8,EP8))*.19,2))} EUR, brutto {euro(round(sum(q*ep for q,ep in zip(Q8,EP8))*1.19,2))} EUR. Die Einheitspreise gelten einschließlich Material, Transport, Montage, Entsorgung und üblicher Nebenarbeiten.",
        ("h","2. Beschaffenheit"),
        "Die Dampfsperre wird auf der freigegebenen Holzschalung verklebt. Die Dämmung wird als Gefälledämmung mit im Mittel 180 mm Dicke nach Plan geliefert. Die zweilagige bituminöse Abdichtung umfasst vollflächige Unterlage und Oberlage. Randanschlüsse erhalten 150 mm Höhe über fertiger Dachoberfläche, ausgenommen der gesondert gezeichnete Türanschluss T-03. Acht Abläufe und zwei Notüberläufe werden an den gezeichneten Stellen eingebaut. Dichtheitskontrolle, Produktunterlagen und Revisionsskizze gehören zur Position 1.08.",
        ("h","3. Termine, Änderung und Zahlung"),
        "Ausführung vom 29.06. bis 17.07.2026. Der Unternehmer meldet geplante Verdeckung kontrollbedürftiger Anschlüsse zwei Arbeitstage vorher. Änderungen und Zusatzpreise vereinbart allein Tessa Lühr. Abschläge werden nach nachgewiesener Leistung gezahlt. Die Schlussrechnung ist nach Abnahme und prüfbarer Abrechnung binnen 30 Tagen zahlbar. Eine pauschale Vertragssicherheit ist nicht vereinbart. Mängelansprüche für die Bauwerksleistung verjähren nach vereinbarter Frist von fünf Jahren ab Abnahme. Abnahme erfolgt in gemeinsamem Termin durch die Geschäftsführerin.",
        "Für die Bauherrin: gez. Tessa Lühr. Für Allerfirst: gez. Nils Gerdes mit schriftlich vorgelegter Abschlussvollmacht der Geschäftsführung. Ausfertigungen an beide Parteien und Rundbogen Architektur."
    ])
    a(3,"Detail_D21_C.png","2026-07-03","vplan","dach","D-21 Revision C, Dach und Türanschluss T-03",[])
    a(4,"Planversand.eml","2026-07-08","vplan","dach","D-21 C jetzt im Verteiler Dach",[
        "Anbei erhalten Sie D-21 Revision C vom 03.07.2026. Der Plan lag seit Freitag im Ordner Fachkoordination; der Versand an Ihre Bauleitung wurde versehentlich nicht ausgelöst. Ihr Ausdruck B im Baustellencontainer ist damit zu ersetzen. Die Änderung betrifft die Aufkantung neben T-03 und den seitlichen Anschluss der Rinne, nicht die ebene Dachfläche.",
        "Herr Gerdes teilte heute telefonisch mit, die Anschlussbahn sei bereits am 06.07.2026 nach B eingebaut und am 07.07. durch das Türelement verdeckt worden. Die Begehung am 09.07.2026 um 08:30 Uhr ist mit Frau Lühr abgestimmt. Eine Preiszusage enthält diese Nachricht nicht. Bitte bewahren Sie Ihre Einbaufotos im ursprünglichen Format auf."
    ],(3,))
    a(5,"Bautagebuch.csv","2026-07-17","vplan","kita","Bautagebuch Dach, Auszug Tagesblätter 31 bis 38",[
        ["Tag","Zeitraum","Beobachter","Ort","Tätigkeit und Befund","Quelle"],
        ["2026-06-29","08:10-09:00","Elisa Wenke","Dach gesamt","Vier Monteure; Schutzgerüst vorhanden; Schalung trocken an zugänglichen Stellen. Messprotokoll Holzfeuchte liegt beim Tragwerksbüro.","Ortsnotiz 31"],
        ["2026-07-01","14:00-14:45","Elisa Wenke","Dach West","Dampfsperre West offen sichtbar; Nahtprobe durch Polier vorgeführt. Ostbereich nicht betreten.","Ortsnotiz 32"],
        ["2026-07-03","11:00","Nils Gerdes","Dach","Dämmung eingebaut; Material gemäß Lieferschein. Rundbogen nicht vor Ort. Wetter laut Polier trocken.","Telefonmeldung 33"],
        ["2026-07-06","16:20","Nils Gerdes","T-03","Anschluss nach Ausdruck D-21 B eingebaut; am 07.07. kommt Türfirma. Zwei Fotos auf Telefon des Poliers.","E-Mail in Projektpostfach 34"],
        ["2026-07-08","10:35","Elisa Wenke","Büro","D-21 C an Allerfirst versandt. Einbauzustand T-03 nicht selbst gesehen.","Planversand 35"],
        ["2026-07-09","08:30-09:35","Wenke und Gerdes","T-03 und Nordrand","Anschluss hinter Türfuß nicht sichtbar. Sichtbare seitliche Höhe 110 mm gemessen; Sollangabe C daneben 150 mm.","Kontrollprotokoll 36"],
        ["2026-07-13","09:00-09:40","Elisa Wenke","Nordrand","Zusatzanschluss sechs Meter offen sichtbar. Gerdes nennt zwei Tage Wartezeit; Personalnachweis noch nicht vorgelegt.","Ortsnotiz 37"],
        ["2026-07-17","13:15-14:00","Wenke und Gerdes","Dach","Gemeinsame Längenaufnahme; ebene Dachmaße 18 mal 12 m. Zwei Laubfänge liegen noch lose neben Ablauf 7 und 8.","Aufmaß 38"],
    ])
    a(6,"Kontrollprotokoll.docx","2026-07-09","vplan","kita","Kontrolltermin T-03, 09.07.2026",[
        "Am 09.07.2026 waren Elisa Wenke, Nils Gerdes und Tessa Lühr von 08:30 bis 09:35 Uhr an Tür T-03 anwesend. Die Wetterlage war trocken, etwa 19 Grad Celsius nach dem Baustellenthermometer. Besichtigt wurden der sichtbare seitliche Dachanschluss, die Rinne und der nördliche Rand; die Fläche hinter dem montierten Türfuß war nicht zugänglich.",
        ("h","1. Beobachtung"),
        "Wenke maß links neben T-03 mit Stahlmaßstab 110 mm zwischen fertiger Dachoberfläche und oberem sichtbaren Bahnrand. Der zugehörige Punkt ist im Plan D-21 C mit 150 mm bezeichnet. Gerdes legte einen Ausdruck B vor, der an dieser Stelle 100 mm zeigt. Das Foto auf seinem Telefon vom 06.07. zeigt den Anschluss vor der Türmontage, jedoch keinen Maßstab. Eine Kopie wurde beim Termin nicht übergeben.",
        ("h","2. Erklärungen"),
        "Gerdes erklärte, C erst am 08.07. erhalten zu haben. Lühr bestätigte, dass die Türfirma den Fuß bereits am 07.07. geschlossen hatte. Wenke bat um einen gemeinsam abgestimmten Zugang zur verdeckten Stelle. Eine Öffnung wurde während des Termins nicht vorgenommen. Über die technische Erfüllung an der verdeckten Stelle wurde keine Feststellung getroffen.",
        ("h","3. Weiterer Ablauf"),
        "Lühr sagte einen Termin mit der Türfirma für den 14.07. zu. Der nördliche Zusatzanschluss kann vorher offen kontrolliert werden. Eine Preisvereinbarung wurde hier nicht getroffen. Das Protokoll wurde am 10.07. an alle drei Teilnehmer versandt; Gerdes bestätigte am 11.07. die Maßangabe, nicht aber eine Kostenübernahme. gez. Elisa Wenke."
    ])
    a(7,"Gemeinsames_Aufmass.pdf","2026-07-17","vplan","dach","Gemeinsames Aufmaß Dach, Blatt A-04",[
        "Elisa Wenke und Nils Gerdes haben am 17.07.2026 von 13:15 bis 14:00 Uhr die nachstehende Geometrie gemeinsam aufgenommen. Bezug sind Dachgrundriss D-21 und die Vertragspositionen vom 10.04.2026. Gemessen wurde mit Bandmaß, die Dachfläche wurde nicht auf die geneigte Oberfläche umgerechnet.",
        ("t",[["Bereich","Messansatz","Ergebnis","Zuordnung"],["Dachfläche","18,00 m mal 12,00 m","216,00 m²","1.02 bis 1.04"],["Randanschluss","18 + 18 + 12 + 12 + 4 m","64,00 m","1.05"],["Zusatzanschluss Nord","6,00 m","6,00 m","Angebot N1"],["Abläufe","8 Stück vorhanden","8 St","1.06"],["Notüberläufe","2 Stück vorhanden","2 St","1.07"]],[100,145,95,155]),
        "Die vier zusätzlichen Meter in Position 1.05 betreffen die beiden Wangen am Türvorbau und sind im Vertrag ausdrücklich enthalten. Die sechs Meter des Zusatzanschlusses Nord gehören nicht zur ebenen Dachfläche und werden nicht zu deren 216 m² addiert. Die sichtbaren Ränder wurden gemessen; der verdeckte Türfuß T-03 wurde nicht geöffnet.",
        "Die Unterschriften bestätigen ausschließlich die gemeinsame Aufnahme der aufgeführten Maße und Stückzahlen. Vergütungsgrund, Preis und technische Vertragsgerechtigkeit sind damit nicht anerkannt. Eine Abnahme wird nicht erklärt. Die Prüfung der Dichtheit und des verdeckten Anschlusses ist nicht Gegenstand dieses Aufmaßblatts.",
        "Aufgenommen und am 17.07.2026 gegengezeichnet: gez. Elisa Wenke; gez. Nils Gerdes. Planbezug: D-21 C, Maßketten Dachrechteck und Nordanschluss."
    ])
    claimq=[1,222,222,222,64,8,2,1]; main=sum(q*ep for q,ep in zip(claimq,EP8)); net=main+900+980
    a(8,"Schlussrechnung_AD26081.pdf","2026-08-03","dach","kita","Schlussrechnung AD-26081, Leistungszeit Juni/Juli 2026",[
        "Wir rechnen die Dacharbeiten am Erweiterungsbau Kita Mühlenwiese abschließend ab. Die Mengen beruhen auf unseren Ausführungsaufzeichnungen und dem gemeinsamen Aufmaß A-04. Unsere Zusatzflächen an der Nordseite werden in den Flächenpositionen mitgeführt. Rechnungszugang per E-Mail am 03.08.2026.",
        ("t",[["Pos.","Leistung","Menge","EP EUR","GP EUR"]]+[[pos,t,f"{q} {unit}",euro(ep),euro(q*ep)] for (pos,t,unit),q,ep in zip(LV8,claimq,EP8)]+[["N1","Zusatzanschluss/Rüsten","1 psch","900,00","900,00"],["N2","Wartezeit","2 Tage","490,00","980,00"]],[40,180,65,95,115]),
        f"Gesamtleistung netto {euro(net)} EUR. Umsatzsteuer 19 Prozent {euro(round(net*.19,2))} EUR. Gesamtleistung brutto {euro(round(net*1.19,2))} EUR. Erhaltene Abschläge: 17.850,00 EUR am 19.06.2026 und 11.900,00 EUR am 17.07.2026, zusammen 29.750,00 EUR brutto. Verbleibender Zahlbetrag {euro(round(net*1.19,2)-29750)} EUR.",
        "Die Wartezeit N2 betrifft den 09. und 10.07.2026. Wir setzen dafür 490,00 EUR je Tag für Mannschaft und Geräte an. Ein gesonderter Stundenbericht ist dieser Rechnung nicht beigefügt. Den Zusatzanschluss N1 haben wir zu dem per E-Mail bestätigten Preis ausgeführt. Die Beseitigung der im Abnahmeprotokoll benannten Punkte ist für den 07.08. vorgesehen.",
        "Zahlbar gemäß Vertrag nach Abnahme und prüfbarer Abrechnung. Zahlungsweg ist das bei der Auftraggeberin hinterlegte Kreditorenkonto AD-12; dieses Dokument enthält keine neue Bankverbindung. Freistellungs- und steuerliche Stammdaten werden im Kreditorenstamm geführt."
    ])
    a(9,"Nachtragsangebot_N1_N2.pdf","2026-07-10","dach","kita","Nachtragsangebot zum Nordanschluss und Wartezeit",[
        "Nach dem am 08.07. zugegangenen Detail C bieten wir den ergänzten nördlichen Anschluss an. Der Anschluss wird auf sechs Metern mit zusätzlichem Anschlussstreifen und höherer Aufkantung hergestellt. Der vorhandene Aufbau wird örtlich geöffnet und wieder geschlossen. Die Ausführung kann am 13.07. erfolgen.",
        ("t",[["Teil","Ansatz","Netto EUR"],["N1.1 Anschlussstreifen","6 m mal 120,00 EUR","720,00"],["N1.2 örtliches Rüsten","1 psch","180,00"],["N1 zusammen","","900,00"],["N2 Wartezeit 09./10.07.","2 Tage mal 490,00 EUR","980,00"],["Angebot gesamt netto","","1.880,00"],["19 Prozent Umsatzsteuer","","357,20"],["Angebot gesamt brutto","","2.237,20"]],[190,180,125]),
        "N1 ersetzt keine der vereinbarten 216 m² Dachfläche und betrifft nicht den verdeckten Türfuß. Für N2 gehen wir von zwei nicht produktiven Tagen aus. Die Mannschaft war auf der Baustelle gemeldet; einen Stundenbeleg reichen wir nach. Andere Arbeiten wurden nach Mitteilung des Poliers teilweise fortgesetzt. Wir bitten um Bestätigung des Leistungsumfangs und der Preise vor Beginn von N1.",
        "Das Angebot ist bis 14.07.2026 bindend. Nur die Geschäftsführung oder deren bevollmächtigte Vertreter können unsere Preise ändern. Eine technische Freigabe durch die Architektin ersetzt die Beauftragung durch die Bauherrin nicht."
    ])
    a(10,"Teilbeauftragung.eml","2026-07-11","kita","dach","N1 bestätigt; Wartezeit nicht bestätigt",[
        "Ich beauftrage den angebotenen Zusatzanschluss N1.1 und N1.2 zu insgesamt 900,00 EUR netto zuzüglich 171,00 EUR Umsatzsteuer. Ausführung wie beschrieben am 13.07.2026; Rundbogen wird vor dem Schließen informiert. Der verdeckte Anschluss T-03 ist nicht Gegenstand dieser Beauftragung.",
        "Die 980,00 EUR Wartezeit N2 werden nicht beauftragt oder anerkannt. Bitte übermitteln Sie die angekündigten Personal- und Gerätestunden mit Darstellung der an diesen Tagen ausgeführten übrigen Tätigkeiten. Eine abschließende Beurteilung ist damit nicht verbunden. Die übrigen Vertragsbedingungen bleiben unverändert."
    ],(9,))
    a(11,"Zahlungsjournal.csv","2026-08-17","kita","vplan","Kassenjournal Projekt VM-26-D",[
        ["Buchungstag","Beleg","Kreditor","Soll_brutto_EUR","Ist_Zahlung_EUR","Status","Verwendungszweck"],
        ["2026-06-19","AD-26042","AD-12","17850.00","17850.00","ausgeführt","Abschlag 1 Dach"],
        ["2026-07-17","AD-26066","AD-12","11900.00","11900.00","ausgeführt","Abschlag 2 Dach"],
        ["2026-08-03","AD-26081","AD-12",f"{round(net*1.19,2)-29750:.2f}","0.00","zur Prüfung","Schlussrechnung Dach"],
        ["2026-08-17","SR-Plan-06","RA-08","8330.00","8330.00","ausgeführt","Planung Abschlag; nicht Dachvertrag"],
    ])
    a(12,"Kostenjournal.xlsx","2026-08-17","kita","vplan","Projektkostenjournal Erweiterung",[])
    a(13,"Balkenterminplan.png","2026-07-14","vplan","kita","Terminplan TP-06, Fortschreibung 14.07.2026",[])
    a(14,"Abnahme_Dach.pdf","2026-07-30","kita","dach","Abnahmeprotokoll Dachvertrag VM-26-D",[
        "Am 30.07.2026 um 10:00 Uhr fand die Abnahmebegehung der Dacharbeiten mit Tessa Lühr, Elisa Wenke und Nils Gerdes statt. Die Erklärung betrifft den Dachvertrag vom 10.04.2026 einschließlich N1 vom 11.07.2026, nicht die Tür-, Elektro- oder Tragwerksleistungen. Frau Lühr erklärte um 11:15 Uhr die Abnahme unter den nachstehenden Vorbehalten.",
        ("h","1. Festgestellte Punkte"),
        "M-01: Am seitlichen Anschluss neben T-03 ist weiterhin eine sichtbare Höhe von etwa 110 mm dokumentiert. Der Bereich hinter dem Türfuß ist verdeckt; eine gemeinsame Öffnung hat noch nicht stattgefunden. Die Bauherrin behält ihre Rechte wegen der abweichenden Anschlussausführung vor. Allerfirst soll mit der Türfirma einen Kontroll- und Beseitigungstermin bis 07.08. abstimmen. Gerdes bestreitet eine alleinige Verantwortung wegen des verspäteten Planversands.",
        "M-02: Die Laubfangkörbe der Abläufe 7 und 8 sitzen lose und sind gegen Herausheben zu befestigen. Die Bauherrin behält ihre Rechte wegen dieses Mangels vor. Beseitigung bis 07.08.2026. Gerdes bestätigt die Ausführung ohne Zusatzpreis.",
        ("h","2. Unterlagen und Erklärungen"),
        "Eine Dichtheitsprüfaufzeichnung für die ebenen Flächen wurde vorgelegt. Sie enthält keine Aussage zum verdeckten Türfuß. Die Revisionsskizze zum Anschluss T-03 fehlt. Eine Zahlung oder Anerkennung der Schlussrechnung wird nicht erklärt. Ein Vertragsstrafenvorbehalt wird nicht aufgenommen, weil dieser Vertrag keine Vertragsstrafe enthält.",
        "Die Bauherrin nimmt mit den ausdrücklich genannten Mängelvorbehalten ab. Die Teilnehmer unterschreiben für ihre jeweiligen Erklärungen: gez. Tessa Lühr, gez. Elisa Wenke, gez. Nils Gerdes. Gerdes erklärt mit seiner Unterschrift keine Zustimmung zu der Ursachenbewertung von M-01."
    ])
    a(15,"Fertigmeldung.eml","2026-08-08","dach","vplan","Abnahmepunkte erledigt",[
        "Unser Monteur war gestern auf dem Dach. Die beiden Laubfangkörbe sind jetzt befestigt. Am Anschluss T-03 hat er außen eine zusätzliche Bahn aufgebracht; nach seiner Meldung ist der Punkt damit erledigt.",
        "Die Türfirma war nicht vor Ort. Eine Öffnung des Türfußes wurde deshalb nicht vorgenommen. Eine neue Revisionsskizze wird nach dem Urlaub unseres Zeichners geliefert. Bitte berücksichtigen Sie die Erledigung bei der Zahlung unserer Schlussrechnung. Fotos befinden sich auf dem Baustellentelefon, sie sind dieser Nachricht nicht beigefügt."
    ])
    a(16,"Nachkontrolle.pdf","2026-08-11","vplan","kita","Nachkontrolle M-01 und M-02",[
        "Elisa Wenke führte am 11.08.2026 von 09:00 bis 09:40 Uhr mit Hausmeister Levin Bahr eine Sichtkontrolle der im Abnahmeprotokoll bezeichneten Stellen durch. Allerfirst war trotz Einladung nicht vertreten. Das Dach war trocken und zugänglich; der Türfuß blieb geschlossen.",
        ("h","1. Tatsächliche Feststellungen"),
        "M-02: Beide Laubfangkörbe ließen sich bei Handprüfung nicht mehr aus ihrer vorgesehenen Halterung anheben. Die Befestigung war sichtbar. M-01: Seitlich neben T-03 ist eine zusätzliche obere Bahn zu sehen. Der untere Anschluss hinter dem Türfuß und die Verbindung zur Rinne waren nicht sichtbar. Eine Funktionsprüfung bei Wasserbelastung wurde nicht durchgeführt.",
        ("h","2. Dokumentation"),
        "Levin Bahr bestätigte, dass seit 08.08. keine Türfirma tätig war. Eine Öffnungserlaubnis oder ein abgestimmter Termin lagen bei dieser Kontrolle nicht vor. Der Bericht belegt daher nur den sichtbaren Zustand und die Handprüfung der Laubfänge. Er enthält keine vollständige technische Bestätigung des Anschlusses T-03.",
        "Die Revisionsskizze T-03 lag auch am 11.08. nicht vor. Versand dieses Berichts an Bauherrin und Allerfirst am 12.08.2026. gez. Elisa Wenke."
    ])
    a(17,"Objektuebergabe.docx","2026-08-14","vplan","kita","Übergabe Erweiterung Kita Mühlenwiese",[
        "Am 14.08.2026 um 14:00 Uhr wurden die Räume G1, G2, Garderobe und Sanitärbereich des Erweiterungsbaus an Tessa Lühr und Kita-Leitung Marit Heese übergeben. Die Schlüsselübergabe umfasst zwölf Raumschlüssel und zwei Zugangstransponder. Sie ersetzt keine noch fehlende Gewerkeabnahme und keine öffentlich-rechtliche oder technische Nutzungsfreigabe.",
        ("t",[["Unterlage/Einweisung","Stand","Empfang"],["Revisionsgrundriss A-20","12.08.2026","Papier und PDF übergeben"],["Elektroprüfung E-09","10.08.2026","PDF übergeben"],["Türbedienung T1 bis T4","14.08.2026","Heese und Bahr eingewiesen"],["Dachdokumentation","Teilmappe 08.08.2026","Produktdaten vorhanden; T-03-Skizze fehlt"],["Heizungsbedienung","12.08.2026","Bahr eingewiesen"]],[180,100,215]),
        "Die Dachabnahme fand am 30.07.2026 statt. Rohbau wurde am 18.06.2026, Fenster/Türen am 05.08.2026 und Elektro am 10.08.2026 gesondert abgenommen; die Originalprotokolle der übrigen Gewerke verwahrt die Bauherrin. Die Übergabeakte enthält diese drei Protokolle nicht. Der Status des Dachpunkts M-01 wird durch die Schlüsselübergabe nicht verändert.",
        "Hausmeister Bahr verwahrt die Wartungsunterlagen im Technikraum. Die Betreiberin übernimmt die turnusmäßige Sichtkontrolle der zugänglichen Abläufe nach der übergebenen Pflegeanleitung. Die Frage der Anschlussausführung T-03 bleibt davon unberührt. Empfang bestätigt: gez. Tessa Lühr, gez. Marit Heese. Übergeben: gez. Elisa Wenke."
    ])
    a(18,"Fristenuebernahme.txt","2026-08-14","kita","vplan","Kassen- und Objektverwaltung, Fristenübernahme",[
        "Mühlenwiese Betreuung gGmbH, Kinderbogen 32, 27283 Verden\n14.08.2026 | VM-26 | Bearbeitung: Fenja Rost, Objektverwaltung\nAn Rundbogen Architektur\nBetreff: Übernahme der Vertragsdaten aus unserem Ordner",
        "Sehr geehrte Frau Wenke,\nfür unsere Terminverwaltung haben wir die folgenden Angaben aus den unterschriebenen Gewerkeprotokollen erfasst. Es handelt sich um Vertragsdaten, nicht um eine Bestätigung späterer Änderungen des Fristverlaufs.",
        "Dach Allerfirst: Vertrag 10.04.2026; fünf Jahre ab Abnahme; Abnahme 30.07.2026; Vermerk Ende 30.07.2031. Rohbau Haenel: Vertrag 12.02.2026; fünf Jahre; Abnahme 18.06.2026; Vermerk Ende 18.06.2031. Fenster/Türen Lichtkante: Vertrag 23.03.2026; fünf Jahre; Abnahme 05.08.2026; Vermerk Ende 05.08.2031. Elektro Wulken: Vertrag 25.03.2026; fünf Jahre; Abnahme 10.08.2026; Vermerk Ende 10.08.2031.",
        "Die drei letztgenannten Abnahmeoriginale verbleiben im Schrank der Geschäftsführung. Der Eintrag für Rundbogen ist noch leer, weil eine gesonderte Abnahme der Architektenleistung nicht protokolliert wurde. Die Schlüsselübergabe vom heutigen Tag haben wir separat erfasst. Zu den beiden Dachmängeln führen wir weiterhin die Kennungen M-01 und M-02.",
        "Mit freundlichen Grüßen\ngez. Fenja Rost\nObjektverwaltung\nAnlagen: keine"
    ])
    a(19,"Fachbeitrag_Pruefung.pdf","2026-08-10","technik","kita","Fachbericht E-09, Elektroinstallation Erweiterung",[
        "Sven Rölke begleitete am 10.08.2026 die Funktionsprüfung der Elektroinstallation im Erweiterungsbau. Die ausführende Elektro Wulken GmbH legte ihre Messprotokolle vor. Betrachtet wurden Beleuchtung, Steckdosenkreise, Auslösung der Fehlerstromschutzeinrichtungen und die Schnittstelle zur Türsteuerung.",
        ("h","1. Prüfumfang"),
        "Die Messprotokolle EW-142 bis EW-148 sind der Betreiberakte beigefügt. Die geprüften Stromkreise wurden durch den ausführenden Elektromeister freigegeben. Die Funktionsprüfung der Türsteuerung betraf die elektrische Ansteuerung und Rückmeldung, nicht die Abdichtung oder die mechanische Befestigung des Türfußes. Bei Ausfall der Versorgung wurde die vorgesehene Rückmeldung ausgelöst.",
        ("h","2. Abgrenzung"),
        "Eine Prüfung des Dachanschlusses T-03, des Tragwerks oder der öffentlich-rechtlichen Bauleiterpflichten ist nicht Gegenstand dieses Berichts. Die Mitwirkung an der Elektroabnahme ersetzt keine Bauabnahme des Dachloses. Die Betreiberunterweisung wurde mit Levin Bahr durchgeführt und von ihm auf Blatt E-09-2 bestätigt.",
        "Die detaillierten elektrischen Messwerte verbleiben in der Betreiberakte; aus dieser Zusammenfassung sollen keine Ersatzmesswerte abgeleitet werden. gez. Sven Rölke. Empfang am 10.08.2026 bestätigt durch Tessa Lühr."
    ])
    a(20,"Bauherrin_Schlussstand.eml","2026-08-18","kita","vplan","Dachrechnung und noch offene Anschlussstelle",[
        "Unsere Kasse hat die Dachschlussrechnung noch nicht bezahlt. Es gingen bislang nur die beiden Abschläge über zusammen 29.750,00 EUR ab. Die Rechnung nennt 222 m²; im Aufmaß stehen 216 m². Herr Gerdes beruft sich auf die zusätzlichen sechs Meter am Nordrand.",
        "Die Türfirma kann am 24.08.2026 um 08:00 Uhr einen Monteur stellen. Für die Öffnung und den späteren Wiederverschluss wurde noch kein Preis vereinbart. Zur Höhe möglicher Beseitigungskosten am Anschluss T-03 liegt uns keine belastbare Angabe vor. Die zusätzliche Bahn aus der Fertigmeldung kann ich als Bauherrin nicht technisch beurteilen.",
        "Die Betreiberin möchte am 31.08. die neuen Gruppenräume belegen. Die offenen Unterlagen bleiben im Übergabeverzeichnis vermerkt. Ich habe weder auf Mängelrechte verzichtet noch eine zusätzliche Wartezeitvergütung bestätigt."
    ])


def define9():
    a=lambda n,f,dt,i,r,t,c,att=(): add(9,n,f,dt,i,r,t,c,att)
    a(1,"Planervertrag.docx","2020-11-09","stadt","uplan","Planervertrag Rathaus Westflügel, Objektbetreuung",[
        "Die Stadt Uelzen, vertreten durch die Projektstelle Rathaus West, beauftragt Büro Raumkante mit Gebäudeplanung und Objektüberwachung der Dach- und Fassadensanierung am Verwaltungsbogen 37. Die Gebäudeleistungsphasen 1 bis 9 werden vereinbart; die LPH 9 wird nach Abschluss der Ausführung ohne zusätzlichen Abruf ausgeführt. Die Fachplanung der Gebäudetechnik bleibt gesondert.",
        ("h","1. Betreuungsumfang"),
        "Raumkante bewertet die innerhalb der Anspruchsfristen festgestellten Mängel einschließlich erforderlicher Begehungen, längstens bis fünf Jahre seit der jeweiligen Abnahme der betroffenen Bauleistung. Es organisiert die Objektbegehung vor Ablauf der Fristen gegenüber den ausführenden Unternehmen und wirkt bei der Freigabe der vereinbarten Sicherheiten mit. Eine Einschränkung gesetzlicher Ansprüche der Stadt gegen andere Beteiligte wird mit dieser Leistungsbeschreibung nicht vereinbart.",
        ("h","2. Zusätzliche Aufgaben"),
        "Die fortlaufende Überwachung der Beseitigung erst nach Abnahme festgestellter Mängel wird nicht mit der Grundbetreuung beauftragt. Ein solcher Auftrag ist mit Umfang und Vergütung gesondert zu vereinbaren. Die Überwachung der Beseitigung bereits bei Bauabnahme festgestellter Mängel bleibt Bestandteil des vereinbarten Überwachungsauftrags. Regelmäßige Gebäudeverwaltung, Wartungsausführung und jährliche Rundgänge ohne konkreten Anlass sind nicht enthalten.",
        ("h","3. Honorar und Erklärungsbefugnis"),
        "Das Gesamthonorar beträgt 62.000,00 EUR netto zuzüglich Umsatzsteuer. Für den ausdrücklich beschriebenen Betreuungsumfang sind davon 2.400,00 EUR netto vereinbart. Dora Sell darf Befunde und Entscheidungsvorschläge ausarbeiten; Verzicht, Vergleich, Sicherheitenfreigabe und kostenpflichtige Zusatzbeauftragung werden ausschließlich durch die vertretungsberechtigte Stadtverwaltung erklärt. Fristkritische Sachverhalte sind unverzüglich Henrike Falk mitzuteilen.",
        "Für die Stadt: gez. Henrike Falk aufgrund interner Abschlussvollmacht. Für Büro Raumkante: gez. Dora Sell. Beide Parteien bestätigen den vorstehenden Vertragsinhalt am 09.11.2020."
    ])
    a(2,"Dachvertrag.pdf","2021-03-12","stadt","udach","Dachbauvertrag Westflügel UR-21-D",[
        "Die Stadt beauftragt Ilmenaudach Bau GmbH mit der Erneuerung des 144 m² großen Flachdachs auf dem Westflügel. Der Vertrag umfasst Abbruch des alten Aufbaus, Dampfsperre, Gefälledämmung, zweilagige Abdichtung, sämtliche Randanschlüsse sowie vier Abläufe. Preisgrundlage ist ein Pauschalpreis von 78.000,00 EUR netto zuzüglich 14.820,00 EUR Umsatzsteuer, zusammen 92.820,00 EUR.",
        ("h","1. Beschaffenheit und Unterlagen"),
        "Es gelten Plan R-12 Revision D vom 08.03.2021 und Anschlussdetail R-15 vom selben Tag. Randanschlüsse sind mit 150 mm Höhe über fertiger Dachoberfläche auszuführen. Die zwei vorhandenen Lüftungsdurchführungen werden in den Aufbau eingebunden. Spätere Montagen fremder Anlagen sind nicht Bestandteil des Auftrags. Die Abrechnung enthält eine Revisionszeichnung und eine dokumentierte Dichtheitsprüfung vor Abnahme.",
        ("h","2. Abnahme und Mängelansprüche"),
        "Es gelten die werkvertraglichen Regelungen des BGB. VOB/B wird nicht vereinbart. Die Parteien vereinbaren für die Bauwerksleistung fünf Jahre für Mängelansprüche ab Abnahme. Die Abnahme erfolgt förmlich in einem gemeinsamen Termin. Eine Schlüsselübergabe oder bloße Inbetriebnahme anderer Gewerke ersetzt diesen Termin nicht.",
        ("h","3. Sicherung"),
        "Die Stadt darf von der Schlusszahlung fünf Prozent der Bruttoabrechnung als Mängelsicherheit einbehalten. Bei unveränderter Abrechnung sind dies 4.641,00 EUR. Eine Ablösung durch gleichwertige Bürgschaft ist möglich, wurde bislang jedoch nicht erklärt. Nach Ablauf der vereinbarten Mängelfrist ist der nicht mehr zur Sicherung bestehender Ansprüche aus diesem Dachvertrag benötigte Betrag freizugeben. Für bekannte offene Ansprüche ist die benötigte Höhe nachvollziehbar zu bestimmen; eine vollständige Sperre ohne Bezug zum Dachvertrag ist nicht vereinbart.",
        "gez. Henrike Falk für die Stadt; gez. Björn Eilers für Ilmenaudach Bau GmbH. Uelzen, 12.03.2021."
    ])
    a(3,"Abnahme_Dach.pdf","2021-10-15","stadt","udach","Förmliche Abnahme Dach Westflügel",[
        "Am 15.10.2021 nahmen Henrike Falk, Dora Sell und Björn Eilers die Dacharbeiten am Westflügel gemeinsam in Augenschein. Frau Falk erklärte für die Stadt um 12:10 Uhr die Abnahme des Dachvertrags UR-21-D mit den nachstehenden Vorbehalten. Eine Erklärung zu anderen Gewerken oder zur Architektenleistung wird nicht abgegeben.",
        ("h","1. Feststellungen"),
        "A-01: Am Ablauf Nordost fehlt der Laubfangkorb. Eilers sagt Lieferung und Befestigung bis 29.10.2021 zu. Die Stadt behält ihre Rechte wegen dieses Mangels vor. A-02: Die Zuordnung der Produktchargen in der Dokumentationsmappe fehlt. Ergänzung bis 29.10.2021. Sichtbare Wasserflecken an der Saaldecke wurden beim Termin nicht festgestellt.",
        ("h","2. Abschlussvermerk vom 02.11.2021"),
        "Sell kontrollierte am 02.11.2021 den befestigten Laubfangkorb und nahm die nachgereichte Chargenliste entgegen. Beide im Abnahmeprotokoll genannten Punkte sind in diesem Umfang erledigt. Es fand keine erneute Abnahme des gesamten Dachs statt. Die Pauschalabrechnung blieb bei 92.820,00 EUR brutto; die Stadt führt einen Einbehalt von 4.641,00 EUR.",
        "Für die jeweiligen Erklärungen: gez. Henrike Falk; gez. Dora Sell; gez. Björn Eilers. Der Abschlussvermerk ist eine Ergänzung derselben Abnahmeurkunde, unterzeichnet durch Dora Sell am 02.11.2021."
    ])
    a(4,"Abnahme_Metallbau.pdf","2021-11-12","stadt","umetall","Abnahme Fenster und Außentüren UR-21-M",[
        "Am 12.11.2021 fand die gemeinsame Abnahme der Fenster- und Außentürarbeiten mit Henrike Falk, Dora Sell und Maren Knoll statt. Frau Falk erklärte die Abnahme mit dem unten bezeichneten Vorbehalt. Der zugrunde liegende Vertrag vom 19.04.2021 umfasst Fenster und Türen zu 47.800,00 EUR netto, 9.082,00 EUR Umsatzsteuer und 56.882,00 EUR brutto. Für Mängelansprüche sind fünf Jahre ab Abnahme und ein Einbehalt von fünf Prozent brutto vereinbart.",
        ("h","1. Vorbehalt M-04"),
        "Der äußere Flügel der Hoftür H2 schleift beim Öffnen an der unteren Führung. Knoll sagt Justierung bis 26.11.2021 zu. Die Stadt behält ihre Rechte wegen dieses Mangels vor. Die übrigen bei der Begehung geöffneten Elemente funktionierten ohne Auffälligkeit. Ein Eingriff in die Dachabdichtung war nicht Gegenstand dieser Leistung.",
        ("h","2. Nachlauf"),
        "Auf diesem Original befindet sich kein Abschlussvermerk für M-04. Ein handschriftlicher Zusatz der Hausverwaltung vom 10.01.2022 lautet: ‚H2 öffnet schwer, Zugang meist über Haupteingang.‘ Die Schlusszahlung wurde unter Einbehalt von 2.844,10 EUR ausgeführt. Über eine Freigabe wurde bis zur Ablage dieser Ausfertigung nicht entschieden.",
        "gez. Henrike Falk; gez. Dora Sell; gez. Maren Knoll. Den Zusatz vom 10.01.2022 zeichnete Hausverwalterin Rena Mert mit Namenskürzel RM."
    ])
    a(5,"Betriebsuebergabe.docx","2022-01-10","stadt","uplan","Betriebsübergabe Westflügel",[
        "Die Projektstelle übergab den Westflügel am 10.01.2022 an die Hausverwaltung, vertreten durch Rena Mert. Der Empfang umfasst die Räume W01 bis W08, den Sitzungssaal und den Zugang zum Dach. Sechs Dachzugangsschlüssel sowie die Papiermappen Dach und Metallbau wurden übernommen.",
        ("h","1. Vorhandene Unterlagen"),
        "Die Dachmappe enthält Abnahme vom 15.10.2021 mit Ergänzung vom 02.11.2021, Produktchargenliste und Bestandsplan R-12 D. Die Metallbaumappe enthält Abnahme vom 12.11.2021; ein Abschlussnachweis für Hoftür H2 fehlt. Die Hausverwaltung übernimmt den Wartungsordner und beauftragt künftig die regelmäßige Reinigung der Dachabläufe.",
        ("h","2. Erklärungen"),
        "Diese Niederschrift dokumentiert die organisatorische Betriebsübergabe. Bereits erklärte Gewerkeabnahmen bleiben unverändert. Eine neue gemeinsame Abnahme aller Bau- oder Architektenleistungen wird nicht erklärt. Die Projektstelle bleibt für Vertrags- und Sicherheitenfragen zuständig.",
        "Rena Mert merkt an, dass H2 weiterhin schwer öffnet und derzeit wenig benutzt wird. Ein Termin mit Heideprofil wurde noch nicht festgelegt. Empfang: gez. Rena Mert. Übergabe: gez. Henrike Falk. Kenntnisnahme: gez. Dora Sell."
    ])
    a(6,"Dachplan_Befundorte.png","2026-09-08","stadt","uplan","R-12 D mit Befundmarkierung Hausverwaltung",[])
    a(7,"Wartung_2025.pdf","2025-10-22","wart","stadt","Wartungsbericht WP-25122, Dach Westflügel",[
        "Leon Tamm war am 22.10.2025 von 08:20 bis 10:10 Uhr mit einer zweiten Servicekraft auf dem Dach Westflügel. Zugang erfolgte über die feste Innentreppe. Es war trocken, die Dachoberfläche stellenweise feucht. Auftrag war Sichtkontrolle und Reinigung der vier Abläufe, nicht eine Bauteilöffnung oder Abdichtungsprüfung mit Wasserbelastung.",
        ("t",[["Stelle","Beobachtung","Arbeit"],["Ablauf NW","Laubauflage am Korb","gereinigt"],["Ablauf NO","Korb befestigt; kleine Zweige","gereinigt"],["Ablauf SW","frei","Sichtkontrolle"],["Ablauf SO","frei","Sichtkontrolle"],["Anschluss Lüftung L2","äußerlich ohne Ablösung","keine Öffnung"]],[105,225,165]),
        "An den zugänglichen Oberflächen wurden keine offenen Nähte erkannt. Die Rückseite der Lüftung L2 war wegen eines abgestellten Gerüstteils nicht vollständig einsehbar. Es wurde kein Wasser in das Gebäude eingebracht. Eine Feuchtemessung im Sitzungssaal war nicht beauftragt.",
        "Die Hausverwaltung erhielt den Bericht am 24.10.2025. Für 2026 ist noch kein Wartungstermin bestätigt. Auftragssumme dieses Termins: 420,00 EUR netto zuzüglich 79,80 EUR Umsatzsteuer. gez. Leon Tamm; Empfang bestätigt: gez. Rena Mert."
    ])
    a(8,"Maengelmeldung.eml","2026-09-08","stadt","uplan","Wasserflecken Sitzungssaal und Tür H2",[
        "Frau Mert hat nach dem Regen am 07.09. im Sitzungssaal zwei Wasserflecken gesehen. Der größere liegt unter der südöstlichen Dachhälfte nahe der Lüftung L2, der kleinere an der Westwand. Sie hat die Stellen im beigefügten Plan mit U-11 und U-12 markiert. Nach ihrer Erinnerung gab es im August einmal einen schwachen Rand; damals wurde keine Meldung geschrieben.",
        "Die Hoftür H2 schleift weiterhin. Frau Mert findet keinen Nachweis, dass Heideprofil den Punkt aus der Abnahme 2021 je abgeschlossen hat. Die Tür wurde seitdem kaum benutzt. Im Saal finden nächste Woche Ausschusssitzungen statt; das Dach ist nur über die Hausverwaltung zugänglich.",
        "Die beiden Sicherungseinbehalte stehen noch in unserer Kasse. Ilmenaudach hatte letzte Woche telefonisch nach Rückzahlung im Oktober gefragt. Ich habe keine Zusage gegeben."
    ],(6,))
    a(9,"Befundblatt_Hausverwaltung.pdf","2026-09-09","stadt","uplan","Hausverwaltung, Befundaufnahme W-Saal",[
        "Rena Mert hielt am 09.09.2026 um 07:45 Uhr den sichtbaren Zustand im Sitzungssaal fest. U-11: unregelmäßiger brauner Rand an der Decke, etwa 62 cm mal 28 cm, gemessen mit Zollstock. U-12: kleiner dunkler Rand an der Westwand, etwa 18 cm breit. Es tropfte bei der Aufnahme nicht. Die Raumtemperaturanzeige zeigte 21 Grad Celsius.",
        ("h","1. Anzeigen und Lage"),
        "Ein ausgeliehenes Handgerät zeigte an U-11 den Wert 78 und am unauffälligen Vergleichspunkt den Wert 32. Das Gerät war auf ‚Baustoff‘ eingestellt; eine Kalibrierung oder Materialzuordnung wurde nicht dokumentiert. Die Anzeigen werden ohne Einheit wiedergegeben. Es wurde kein Material entnommen und keine Decke geöffnet.",
        ("h","2. Dachzugang"),
        "Mert sah vom Dachausstieg aus Laub am nordwestlichen Ablauf. Die südöstliche Fläche wurde nicht betreten, weil ein Geländerabschnitt nach Arbeiten an der Lüftung nicht wieder eingesetzt war. Die Firma für diese Arbeiten war nach dem Hausbuch am 24.08.2026 vor Ort. Ob an der Abdichtung gearbeitet wurde, ist im Hausbuch nicht vermerkt.",
        "Die Maße bezeichnen sichtbare Ränder, nicht die räumliche Ausdehnung feuchten Materials. Die Planmarkierung vom Vortag wurde unverändert verwendet. gez. Rena Mert. Anlagen: keine."
    ])
    a(10,"Unternehmerantwort.eml","2026-09-14","udach","stadt","Besichtigung ohne Anerkennung",[
        "Wir haben Ihre Nachricht vom 10.09.2026 mit den beiden Befundstellen erhalten. Nach unserer Abnahme 2021 gab es bis heute keine Wasserbeanstandung. Bitte übersenden Sie Wartungsnachweise und Angaben zu den Arbeiten an Lüftung L2 im August.",
        "Ohne Anerkennung einer Rechtspflicht können wir am 14.10.2026 um 15:00 Uhr eine Sichtprüfung anbieten. Eine Reparaturzusage und ein Anerkenntnis einer Ursache sind damit nicht verbunden. Unser Monteur benötigt sicheren Zugang zur südöstlichen Dachfläche. Eine Verlängerung der vertraglichen Fristen erklären wir nicht.",
        "Bitte bestätigen Sie außerdem, wann der Sicherungseinbehalt von 4.641,00 EUR ausgezahlt werden soll. Uns liegt bislang keine bezifferte Gegenforderung vor."
    ])
    a(11,"Begehung_September.pdf","2026-09-21","uplan","stadt","Begehungsniederschrift vom 21.09.2026",[
        "Dora Sell und Rena Mert waren am 21.09.2026 von 09:00 bis 10:20 Uhr im Westflügel. Ilmenaudach und Heideprofil waren nicht anwesend. Der Sitzungssaal wurde besichtigt, das Dach nur vom Ausstieg aus. Der fehlende Geländerabschnitt erlaubte keine Begehung der südöstlichen Fläche im vorgesehenen Ablauf.",
        ("h","1. Sichtbarer Befund"),
        "U-11 zeigte einen trockenen braunen Rand in der angegebenen Lage unterhalb des Bereichs L2. U-12 war schwächer erkennbar. Es wurden keine Bauteile geöffnet und keine Materialfeuchten bestimmt. Die Zuordnung der Deckenstellen zu Dachanschlüssen beruht auf Plan R-12, nicht auf einer geöffneten Konstruktion. Ein Zusammenhang mit L2 ist möglich, aber durch diese Begehung nicht nachgewiesen.",
        ("h","2. Tür und Dokumente"),
        "H2 schleifte bei einem Öffnungsversuch an der unteren Führung. Die Stelle entspricht der Beschreibung M-04 aus dem Abnahmeprotokoll. Ein Nachweis früherer Justierung lag nicht vor. Wartungsbericht 2025 wurde gelesen; er enthält keine vollständige Prüfung der Rückseite L2. Für die Augustarbeiten fehlt der Montagebericht.",
        ("h","3. Terminnotiz"),
        "Falk wurde am 21.09. telefonisch über den unvollständigen Dachzugang informiert. Sie nannte den 18.10. als organisatorisch möglichen gemeinsamen Termin. Sell hielt fest, dass Dachabnahme und Metallbauabnahme an verschiedenen Tagen im Jahr 2021 erfolgt sind. Diese Niederschrift enthält keine Vereinbarung mit den Unternehmen über Fristen oder Ansprüche. gez. Dora Sell; Anwesenheit bestätigt: gez. Rena Mert."
    ])
    a(12,"Sicherungen.xlsx","2026-09-24","stadt","uplan","Kassenübersicht Sicherungseinbehalte Westflügel",[])
    a(13,"Sicherungskonto.csv","2026-09-24","stadt","uplan","Kassenkonto 8842, Bewegungen Sicherung Westflügel",[
        ["Datum","Vertrag","Buchung","Soll_EUR","Haben_EUR","Beleg","Text"],
        ["2021-11-18","UR-21-D","Einbehalt","4641.00","0.00","SR-D-21104","5 Prozent von 92820 EUR"],
        ["2021-12-16","UR-21-M","Einbehalt","2844.10","0.00","SR-M-21088","5 Prozent von 56882 EUR"],
        ["2026-09-24","UR-21-D","Freigabe","0.00","0.00","Kassenstand","keine Auszahlungsanordnung"],
        ["2026-09-24","UR-21-M","Freigabe","0.00","0.00","Kassenstand","keine Auszahlungsanordnung"],
    ])
    a(14,"Terminabstimmung.eml","2026-09-23","stadt","uplan","Gemeinsamer Termin im Oktober",[
        "Der Saal ist vom 12. bis 16.10. durch Veranstaltungen belegt. Frau Mert und unser Hausdienst könnten den 18.10.2026 ab 09:00 Uhr ermöglichen. Dieser Tag wäre aus Sicht der Hausverwaltung für Dach und Metallbau gemeinsam praktisch.",
        "Die Wiederherstellung des Geländerabschnitts ist für den 02.10. vorgesehen, aber noch nicht bestätigt. Ein Angebot für Bauteilöffnungen liegt nicht vor. Ilmenaudach bietet seinerseits nur den 14.10. an. Die ursprünglichen Abnahmen waren Dach am 15.10.2021 und Metallbau am 12.11.2021; die Schlüsselübergabe war am 10.01.2022.",
        "Eine Fristvereinbarung oder Beauftragung zusätzlicher Überwachung habe ich bisher nicht unterschrieben. Herr Eilers hat der Terminverschiebung noch nicht zugestimmt."
    ])
    a(15,"Telefonnotiz.txt","2026-09-24","stadt","uplan","Telefonnotiz Augustarbeiten",[
        "Stadt Uelzen, Projektstelle Rathaus West\nVerwaltungsbogen 37, 29525 Uelzen\n24.09.2026, 14:10 bis 14:22 Uhr | UR-21-West\nBearbeitung: Henrike Falk\nGespräch mit Rena Mert, Hausverwaltung",
        "Frau Mert berichtet, dass am 24.08.2026 zwei Mitarbeiter der Lüftung Wartwerk GmbH einen Motor an L2 austauschten. Die Firma nutzte den Dachausstieg und stellte Werkzeug neben dem Anschluss ab. Mert war nicht durchgehend anwesend. Sie kann nicht sagen, ob die Abdichtung betreten, beschädigt oder geöffnet wurde. Der Montagebericht ist noch beim technischen Gebäudemanagement angefordert.",
        "Nach dem Hausbuch trat am 28.08. ein schwacher Deckenrand auf. Mert hielt ihn zunächst für einen alten Fleck. Erst nach dem Regen vom 07.09. wurde der Rand deutlich größer. Eine Reinigung des nordwestlichen Ablaufs im Jahr 2026 ist nicht eingetragen. Für U-12 kann Mert keinen genauen ersten Tag nennen.",
        "Das Gespräch enthielt keine technische Diagnose. Mert sagte zu, den sicheren Dachzugang mit dem Hausdienst zu klären. Falk hat keine Reparatur beauftragt und keine Aussage zur Verantwortlichkeit gegenüber einem Unternehmen abgegeben.\n\ngez. Henrike Falk\nAnlagen: keine"
    ])
    a(16,"Angebot_Zusatzbetreuung.docx","2026-09-25","uplan","stadt","Angebot gesonderte Beseitigungsüberwachung U-11/U-12",[
        "Für den Fall einer beauftragten Mängelbeseitigung an den später gemeldeten Befundstellen U-11 und U-12 bieten wir folgende zusätzliche Leistung an. Unsere bereits vereinbarte fachliche Bewertung, die notwendige Begehung hierzu, die Vorfristbegehung und die Mitwirkung bei der Sicherheitenfreigabe bleiben unverändert Bestandteil des Ausgangsvertrags.",
        ("h","1. Zusätzlicher Leistungsumfang"),
        "Wir koordinieren zwei Kontrolltermine während einer fachlich festgelegten Beseitigungsmaßnahme und einen dokumentierten Abschlusstermin. Vor einer Verdeckung werden die im Beseitigungskonzept bezeichneten Anschlüsse mit der ausführenden Firma zugänglich gehalten. Die Kontrollen werden mit Ort, Datum, Beteiligten und tatsächlich sichtbaren Befunden dokumentiert. Die technische Spezialdiagnostik, Bauteilöffnungen und die Planung eines neuen Dachaufbaus sind nicht enthalten.",
        ("h","2. Vergütung und Voraussetzungen"),
        "Das zusätzliche Honorar beträgt pauschal 1.800,00 EUR netto zuzüglich 342,00 EUR Umsatzsteuer, zusammen 2.142,00 EUR. Voraussetzung sind sicherer Zugang und rechtzeitige Mitteilung der Ausführungstermine. Weitere Termine werden nur nach zusätzlicher Vereinbarung durchgeführt. Eine Befugnis zu Preisvereinbarungen, Abnahmeerklärungen, Verzicht oder Sicherheitenfreigabe wird nicht erteilt.",
        ("h","3. Angebotsstatus"),
        "Dieses Angebot ist bis 09.10.2026 bindend und noch nicht angenommen. Für den alten Abnahmepunkt M-04 der Hoftür wird mit diesem Angebot kein neuer Leistungsumfang berechnet. Über dessen bisherigen Bearbeitungsstand werden die Unterlagen des Ausgangsauftrags gesondert abgeglichen. Die Auftraggeberin erhält diesen vollständigen Text als Vertragsangebot; eine Annahmeerklärung liegt Büro Raumkante derzeit nicht vor."
    ])


def image_base(d, size=(1800,1300)):
    im=Image.new("RGB",size,"white"); draw=ImageDraw.Draw(im)
    draw.text((60,35),ACTORS[d["issuer"]][0],font=screen_font(32,True),fill="#1e3540")
    draw.text((60,86),d["title"],font=screen_font(29,True),fill="black")
    draw.text((60,130),CASES[d["phase"]]["ref"]+" | "+d["date"]+" | "+ACTORS[d["issuer"]][3],font=screen_font(23),fill="#444444")
    draw.rectangle((40,180,size[0]-40,size[1]-110),outline="#777777",width=2)
    draw.text((60,size[1]-84),"Maße haben Vorrang. Ausdruck nicht maßstäblich. Bezug und Revisionsstand im Planfeld.",font=screen_font(21),fill="#333333")
    draw.text((60,size[1]-49),ACTORS[d["issuer"]][1]+" | "+ACTORS[d["issuer"]][2],font=screen_font(20),fill="#444444")
    return im,draw


def label(draw,xy,text,size=23,color="black",bold=False):
    draw.multiline_text(xy,text,font=screen_font(size,bold),fill=color,spacing=7)


def dimension(draw,start,end,text,vertical=False):
    draw.line([start,end],fill="#353535",width=2)
    for x,y in (start,end):
        draw.line((x-7,y-7,x+7,y+7),fill="#353535",width=2)
    x=(start[0]+end[0])/2; y=(start[1]+end[1])/2
    label(draw,(x+12 if vertical else x-45,y-36 if not vertical else y),text,20)


def save_image(im,d):
    metadata=PngImagePlugin.PngInfo(); metadata.add_text("Author",AUTHOR); metadata.add_text("Title",d["title"])
    metadata.add_text("Description",CASES[d["phase"]]["ref"]+" | "+d["date"])
    im.save(doc_path(d),pnginfo=metadata,dpi=(150,150))


def plan7(d):
    im,draw=image_base(d)
    label(draw,(70,204),"F-04 B | Elementübersicht Leseflügel | Revision B: Sonnenschutzzuordnung",24,bold=True)
    # Drei echte Fassadenstreifen mit allen ausgeschriebenen Elementen.
    elevations=[("Südseite",250,["W2"]*6), ("Westseite",510,["W1"]*6+["W2"]*2), ("Ostseite",770,["W1"]*6+["T1"]*2)]
    counts={"W1":0,"W2":0,"T1":0}
    for title,y,elements in elevations:
        label(draw,(75,y),title,22,bold=True)
        draw.line((70,y+185,1720,y+185),fill="#333333",width=4)
        x=230
        for kind in elements:
            counts[kind]+=1; width=150 if kind=="W2" else 105; height=140 if kind=="T1" else 100
            bottom=y+175; top=bottom-height
            draw.rectangle((x,top,x+width,bottom),fill="#e8f1f3",outline="#1d4d60",width=4)
            if kind=="W2":draw.line((x+width*.45,top,x+width*.45,bottom),fill="#1d4d60",width=3)
            else:draw.line((x,top,x+width,bottom),fill="#7895a0",width=2)
            if title=="Südseite":
                for yy in range(top+7,bottom,12):draw.line((x+3,yy,x+width-3,yy),fill="#7d878b",width=2)
            label(draw,(x,top-31),f"{kind}-{counts[kind]:02d}",18)
            x+=width+42
    label(draw,(80,1030),"W1: 12 Stück, 1,20 x 1,50 m\nW2: 8 Stück, 1,80 x 1,50 m\nT1: 2 Stück, 1,30 x 2,40 m",22)
    label(draw,(620,1030),"R1: 6 Raffstores nur Süd\nSchwelle T1: 20 mm\nUw W1/W2: höchstens 0,95",22)
    label(draw,(1200,1030),"Putzanschlüsse: 72 m\n20 Fensterbänke\nElektroanschluss separates Los",22)
    save_image(im,d)


def plan8(d):
    im,draw=image_base(d)
    label(draw,(70,205),"D-21 C | Dachgrundriss und Schnitt T-03 | Änderung 03.07.2026",24,bold=True)
    x,y,w,h=160,340,810,540
    draw.rectangle((x,y,x+w,y+h),fill="#f0f1f1",outline="#233d48",width=8)
    draw.line((x+w/2,y,x+w/2,y+h),fill="#9aabb2",width=2)
    dimension(draw,(x,y-45),(x+w,y-45),"18,00 m")
    dimension(draw,(x-55,y),(x-55,y+h),"12,00 m",True)
    label(draw,(x+265,y+230),"Dachfläche 216,00 m²\n18,00 x 12,00 m",26,bold=True)
    points=[(x+50,y+50),(x+w-50,y+50),(x+50,y+h-50),(x+w-50,y+h-50),(x+240,y+50),(x+570,y+50),(x+240,y+h-50),(x+570,y+h-50)]
    for i,(px,py) in enumerate(points,1):
        draw.ellipse((px-12,py-12,px+12,py+12),fill="#306578");label(draw,(px+15,py-15),f"A{i}",18)
    draw.rectangle((x+335,y+h-10,x+465,y+h+25),fill="#cea33e",outline="black",width=2)
    label(draw,(x+340,y+h+36),"T-03",22,bold=True)
    draw.line((x+120,y-10,x+390,y-10),fill="#a34343",width=10)
    label(draw,(x+180,y-100),"N1: 6,00 m",20,color="#922e2e")
    label(draw,(1080,275),"Schnitt neben T-03",24,bold=True)
    draw.rectangle((1090,660,1630,715),fill="#9b9b9b",outline="black",width=2)
    draw.polygon([(1090,660),(1450,615),(1450,660)],fill="#efdbaa",outline="black")
    draw.line((1090,648,1450,600,1450,450),fill="#2c5369",width=9)
    draw.rectangle((1470,350,1520,720),fill="#e3e5e6",outline="#333333",width=3)
    draw.line((1090,720,1650,720),fill="#333333",width=3)
    dimension(draw,(1570,450),(1570,600),"150 mm",True)
    draw.line((1370,500,1480,500),fill="#a34343",width=3)
    label(draw,(1100,385),"Revision C: 150 mm",21,color="#234b62")
    label(draw,(1090,530),"B: 100 mm, ersetzt",20,color="#922e2e")
    label(draw,(1070,780),"Dampfsperre auf Holzschalung\nGefälledämmung im Mittel 180 mm\nAbdichtung zweilagig\nTürfuß hinter Profil nicht sichtbar",21)
    label(draw,(80,1010),"Randanschluss: 18 + 18 + 12 + 12 + 4 = 64 m. N1 gesondert, keine zusätzliche ebene Dachfläche.\nZwei Notüberläufe an Ost-/Westseite. Dachentwässerung in Fachplan E-14.\nRevision C ändert die seitliche Aufkantung und den Rinnenanschluss an T-03.",23)
    save_image(im,d)


def timeline8(d):
    im,draw=image_base(d)
    label(draw,(70,210),"TP-06 | Basis dunkelgrau, Stand 14.07. blau, noch vorgesehen hellblau",24,bold=True)
    begin=datetime(2026,6,29); x0=440; step=22
    for day in range(0,48,7):
        x=x0+day*step; draw.line((x,285,x,1030),fill="#d3d8da",width=2)
        from datetime import timedelta
        label(draw,(x-18,255),(begin+timedelta(days=day)).strftime("%d.%m."),20)
    tasks=[("Dampfsperre",0,3,0,3),("Dämmung",3,6,3,6),("Abdichtung",6,11,6,11),("Klärung T-03",8,10,9,16),("Zusatzanschluss N1",10,12,14,15),("Dach fertig",16,18,16,18),("Dachabnahme",24,25,31,32),("Übergabe",40,41,46,47)]
    for row,(title,b,e,nb,ne) in enumerate(tasks):
        y=315+row*83;label(draw,(80,y+10),title,24)
        draw.rectangle((x0+b*step,y,x0+e*step,y+17),fill="#666b6d")
        draw.rectangle((x0+nb*step,y+25,x0+ne*step,y+53),fill="#467a90" if nb<=15 else "#b9d4de")
    label(draw,(80,1080),"T-03: verdeckter Bereich ohne bestätigten Kontrolltermin. Dachabnahme für 30.07. vorgesehen.\nÜbergabe 14.08.; Betrieb ab 31.08. noch abhängig von gesonderten Freigaben.",23)
    save_image(im,d)


def plan9(d):
    im,draw=image_base(d)
    label(draw,(70,205),"R-12 D | Bestandsgeometrie 08.03.2021 | Markierungen Rena Mert 08.09.2026",23,bold=True)
    x,y,side=240,355,600
    draw.rectangle((x,y,x+side,y+side),fill="#eeeeed",outline="#263e46",width=8)
    dimension(draw,(x,y-45),(x+side,y-45),"12,00 m")
    dimension(draw,(x-55,y),(x-55,y+side),"12,00 m",True)
    draw.line((x+300,y,x+300,y+side),fill="#8d999f",width=2)
    for title,px,py in [("NW",x+40,y+40),("NO",x+560,y+40),("SW",x+40,y+560),("SO",x+560,y+560)]:
        draw.ellipse((px-14,py-14,px+14,py+14),fill="#376579");label(draw,(px+20,py-15),title,20)
    for title,px,py in [("L1",x+175,y+195),("L2",x+440,y+410)]:
        draw.rectangle((px-35,py-35,px+35,py+35),fill="#c5d5dc",outline="#2e4652",width=3);label(draw,(px-18,py-12),title,20,bold=True)
    draw.ellipse((x+410,y+440,x+515,y+490),outline="#a03d3d",width=5);label(draw,(x+520,y+460),"U-11",22,color="#8a2929")
    draw.ellipse((x-15,y+360,x+45,y+410),outline="#a03d3d",width=5);label(draw,(x-105,y+360),"U-12",22,color="#8a2929")
    draw.rectangle((x+260,y+510,x+335,y+side),fill="#d8c78e",outline="#333333",width=2);label(draw,(x+250,y+620),"Dachausstieg",22)
    draw.line((920,390,920,275),fill="black",width=4);draw.polygon([(920,260),(909,285),(931,285)],fill="black");label(draw,(905,225),"N",24,bold=True)
    label(draw,(1100,340),"Westflügel Rathaus\nDachfläche 144,00 m²\nSitzungssaal unter Südteil",25,bold=True)
    label(draw,(1100,510),"U-11: Deckenrand 62 x 28 cm\nU-12: Wandrand etwa 18 cm\nÜbertragung aus Raumbeobachtung",23,color="#8a2929")
    label(draw,(1100,670),"L2: Lüftungsdurchführung\nArbeiten am 24.08.2026 gemeldet\nGeländerabschnitt Südost fehlt",23)
    label(draw,(1100,820),"Kreise zeigen Befundorte im Raum,\nnicht nachgewiesene Eintrittsstellen.\nKeine Bauteilöffnung durchgeführt.",23)
    label(draw,(80,1090),"Grundlage R-12 D. Lageübertragung durch Hausverwaltung; keine neue Vermessung.\nLegende: blaue Punkte Abläufe; Rechtecke L1/L2 Durchführungen; rot Befunde unterhalb des Dachs.",23)
    save_image(im,d)


DESCRIPTIONS={
    7:"Ein privates Stiftungsprojekt mit drei vollständigen Fensterangeboten, verbindlicher Antwort zur Elementleistung, bedingtem Skonto, abweichender Montagezeit und noch offener Förderentscheidung. Präzises Übungsgebiet: positionsbezogene Angebotswertung, zulässige Aufklärung, Auftraggeber- und Förderbindung sowie Zusammenstellung des Vertragssatzes ohne eigenen Zuschlag.",
    8:"Erweiterung einer Kita mit verspätem Detailversand, konkret protokollierter Ortskontrolle, verdecktem Türanschluss, gemeinsamem geometrischem Aufmaß, kumulativer Schlussrechnung und nur teilweise bestätigtem Nachtrag. Präzises Übungsgebiet: tatsächliche Bauüberwachung mit Erkenntnisgrenzen, Mengen-/Kostenabgleich, Abnahmevorbehalte, Übergabe und Nachkontrolle der Abnahmemängel.",
    9:"Rathaussanierung mit Dach- und Metallbauabnahmen aus 2021, späteren Feuchtebefunden, Wartungsnachweisen, fremdem Eingriff und fehlendem Dachzugang. Präzises Übungsgebiet: fachliche Bewertung späterer Mängel, Begehung vor gewerkeweiser Anspruchsfrist, Sicherungseinbehalte und Abgrenzung einer zusätzlichen Beseitigungsüberwachung.",
}
REVIEW_POINTS={
    7:[("angebote","Alle acht Positionen aus drei vollständigen Angeboten vergleichen; Skonto nicht als unbedingten Nachlass behandeln. Die Erfüllung von Uw und Montagezeit bleibt gesondert zu prüfen."),
       ("verfahren","Stiftungsorganisation und noch nicht bewilligte Förderung aus den Originalen würdigen. Weder automatisch kommunale EU-Vergabe noch vollständige Verhandlungsfreiheit behaupten."),
       ("aufklaerung","BG-78-Antwort klärt die angebotene Beschaffenheit noch nicht vollständig. Neue warme Kante nicht heimlich als unverändertes Angebot übernehmen. Kein selbst ausgeführter Zuschlag."),
       ("vertrag","Vertragsunterlagen nach exaktem Stand mit Antwort vom 07.09. zusammenstellen. Budget, Nettosumme, Bindefrist und Wiedereröffnung voneinander unterscheiden.")],
    8:[("vorort","Tatsächliche Beobachter, Planversand und verdeckte Bereiche unterscheiden. Aus dem nachgereichten Plan oder der Fertigmeldung keine eigene Ortskontrolle oder vollständige Anschlussfreigabe erzeugen."),
       ("mengen","216 m² aus 18 mal 12 geometrisch prüfen; sechs Meter Zusatzanschluss nicht als zusätzliche Fläche übernehmen. Vertragsbasis 36986 EUR netto, bestätigtes N1 900 EUR netto und nicht bestätigtes N2 980 EUR auseinanderhalten."),
       ("rechnung","Unternehmerforderung 39640 EUR netto mit geometrischer Basis abgleichen. Tatsächlich gezahlte 29750 EUR brutto nur einmal abziehen; Rechnungskorrektur und mangelsbezogenen Einbehalt getrennt prüfen."),
       ("abnahme","Abnahme 30.07. mit M-01/M-02 von Übergabe 14.08. trennen. Nachkontrolle bestätigt nur M-02; M-01 bleibt hinter dem Türfuß ungeprüft. Abnahmemängel bleiben LPH-8-Beseitigungsüberwachung."),
       ("kosten","Projektjournal enthält Eingangsrechnungen und Vertragswerte, keine fertige Kostenfeststellung. Gebäudeabschluss nicht allein aus Dachrechnung erklären; Planerleistung, Gebühren und andere Gewerke getrennt führen.")],
    9:[("fristen","Dachabnahme 15.10.2021 und Metallbauabnahme 12.11.2021 unterscheiden. Betriebsübergabe 10.01.2022 nicht als gemeinsamer Beginn. Vorgesehener 18.10.2026 liegt nach vorläufigem Dachfristende."),
       ("ursache","U-11/U-12 nicht allein Dachunternehmer zuweisen; Wartung, L2-Arbeiten und unzugänglicher Bereich bleiben relevante Tatsachen. Gerätewerte 78/32 sind keine Masseprozente."),
       ("leistung","Alten Abnahmemangel M-04 und späteren Feuchtebefund getrennt bearbeiten. Überwachung späterer Beseitigung ist nicht automatisch vom Grundauftrag umfasst; Angebot vom 25.09. noch nicht angenommen."),
       ("sicherung","4641 EUR Dach und 2844.10 EUR Metallbaueinbehalt vertragsbezogen beurteilen. Keine pauschale Sperre aller Sicherheiten und keine automatische Freigabe wegen Zeitablaufs."),
       ("hemmung","Prüfangebot ohne Anerkennung ist weder automatisch Hemmung noch Neubeginn; tatsächliches Verhalten und Vertragslage prüfen. HOAI-Fünfjahresgrenze nicht als allgemeine Verjährungsregel nutzen.")],
}


def README(phase):
    c=CASES[phase]; slug=c["slug"]; files=FILES[phase]
    spec=importlib.util.spec_from_file_location("hoai789_downloads", ROOT/"scripts/inject-gesamt-pdf-section.py")
    downloads=importlib.util.module_from_spec(spec); spec.loader.exec_module(downloads)
    pdf_rel=f"gesamt-pdf/{slug}_gesamt.pdf"
    download_section=downloads.section_block(slug,pdf_rel if (ROOT/"testakten"/slug/pdf_rel).exists() else None,True)
    inventory="\n".join(f"| {n:02d} | [{d['filename']}]({d['filename']}) | {d['date']} | {d['title']} |" for n,d in files.items())
    text=f"""<!-- decimal-headings -->

# {c['title']}

{download_section}

## Umfang und Übungsgebiet

Autor: Klotzkette. Fallkennung: `{slug}`. Aktenstand: {c['date']}. {len(files)} eigenständige Originalunterlagen.

{DESCRIPTIONS[phase]}

## Herkunft

<!-- reserved-example-contacts -->

Sämtliche Personen, Firmen, Anschriften, Grundstücke, Verträge, Verwaltungs- und Zahlungsvorgänge sind erfunden. Die Ortsnamen Melle, Verden und Uelzen sowie die verlinkten amtlichen Normquellen sind real. Auch scheinbar amtliche Unterlagen sind eigens gestaltete Fiktion, keine echten Behördenurkunden. Es werden keine echten Bankdaten, Ausweise, Stempel oder Sicherheitsmerkmale verwendet. Alle Mailadressen verwenden reservierte `.example`-Domains. Pläne und Terminbilder wurden für diese Akte als tatsächliche Rasterzeichnungen hergestellt; sie bilden kein reales Bauvorhaben ab. Die Originale enthalten keine Musterlösung und keine Übungsanweisung.

## Originalinventar

| Nr. | Original | Datum | Inhalt |
| --- | --- | --- | --- |
{inventory}

## Quellen und Grenzen

Am 25.09.2026 wurden [HOAI Anlage 10](https://www.gesetze-im-internet.de/hoai_2013/anlage_10.html), [Paragraf 34 HOAI](https://www.gesetze-im-internet.de/hoai_2013/__34.html) und [Paragraf 650p BGB](https://www.gesetze-im-internet.de/bgb/__650p.html) gezielt gelesen. Die letzten beiden Einzeltexte waren nach Webabruffehlern per direktem amtlichem HTTP-Abruf lesbar. Die [phasebezogenen Quellen](../../bauwirtschaft/references/hoai-{phase}-fachquellen.md) dokumentieren weitere tatsächlich gelesene Normen und Grenzen. Historische Fassungen, Landesrecht, einzelne Sicherungsklauseln und technische Normvolltexte sind damit nicht vollständig geprüft. Es werden keine neuen Rechtsprechungsanker behauptet. Projektmaße und Materialwerte sind erfundene Vertragsvorgaben, keine wörtliche DIN-Wiedergabe.

## Reproduktion und Prüfung

Originalbuilder: `scripts/build-bauwirtschaft-hoai-7-9-akten.py --phase {phase} --qa DIR`. Tabellen werden durch den zugehörigen MJS-Builder mit Artifact Tool erstellt, nicht mit openpyxl. Danach führt `scripts/build-bauwirtschaft-hoai-7-9-qa.py --phase {phase} --qa DIR` die native Office-Neuberechnung, Eingabemutation und das Rendering aus. Die Office-Laufzeit wird über die portablen Helfer beziehungsweise `SOFFICE` gefunden; `DOCX_RENDERER` bezeichnet das vorhandene Dokument-Renderwerkzeug, `DOCX_PYTHON` optional dessen Python-Laufzeit. `scripts/test-bauwirtschaft-hoai-7-9.py` prüft standardmäßig die vorhandenen Originale offline; `--assets DIR` prüft zusätzlich kanonische Exporte. Die interne `rubric.yaml` ist nicht Teil der Archive. Menschliche Ergebnisbewertungen und Live-Modelltests werden nicht als bestanden ausgegeben.
"""
    (ROOT/"testakten"/slug/"README.md").write_text(normalize_decimal_headings(ensure_download_notices(text,case_readme=True)),encoding="utf-8")
    import yaml
    checks=[dict(id=k,check_type="human_review",description=v) for k,v in REVIEW_POINTS[phase]]
    checks.extend(dict(id=f"original-{n:02d}",check_type="file_exists",path=d["filename"],description="Eigenständige Originalunterlage vorhanden.") for n,d in files.items())
    (ROOT/"testakten"/slug/"rubric.yaml").write_text(yaml.safe_dump(dict(name=slug,plugin="bauwirtschaft",description=DESCRIPTIONS[phase],checks=checks),allow_unicode=True,sort_keys=False),encoding="utf-8")


def define_all():
    define7(); define8(); define9()


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--phase",type=int,choices=[7,8,9]); parser.add_argument("--qa",type=Path)
    parser.add_argument("--no-workbooks",action="store_true")
    args=parser.parse_args(); qa=args.qa or Path(tempfile.mkdtemp(prefix="hoai-7-9-build-")); qa.mkdir(parents=True,exist_ok=True)
    define_all(); register_fonts(); phases=[args.phase] if args.phase else [7,8,9]
    for phase in phases:
        folder=ROOT/"testakten"/CASES[phase]["slug"]; folder.mkdir(parents=True,exist_ok=True)
        for d in FILES[phase].values():
            ext=Path(d["filename"]).suffix
            if ext==".pdf":pdf(d)
            elif ext==".docx":docx(d)
            elif ext==".txt":doc_path(d).write_text("\n\n".join(d["content"])+"\n",encoding="utf-8")
            elif ext==".csv":
                with doc_path(d).open("w",encoding="utf-8-sig",newline="") as f:csv.writer(f).writerows(d["content"])
            elif ext==".png":{(7,5):plan7,(8,3):plan8,(8,13):timeline8,(9,6):plan9}[(phase,d["number"])](d)
        for d in FILES[phase].values():
            if d["filename"].endswith(".eml"):eml(d)
        README(phase)
    data=dict(cases={str(n):CASES[n] for n in phases},files={str(n):FILES[n] for n in phases},q7=Q7,ep7=EP7,lv7=LV7,q8=Q8,ep8=EP8,lv8=LV8)
    config=qa/"build-data.json"; config.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding="utf-8")
    if not args.no_workbooks:
        subprocess.run([node_binary(),str(ROOT/"scripts/build-bauwirtschaft-hoai-7-9-workbooks.mjs"),"--data",str(config),"--qa",str(qa)],check=True)
    print(json.dumps(dict(phases=phases,originals=sum(len(FILES[n]) for n in phases),qa=str(qa)),ensure_ascii=False))


if __name__=="__main__":main()
