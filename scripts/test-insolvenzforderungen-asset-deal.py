#!/usr/bin/env python3
"""Prüft Bestandsdaten und die beiden ergänzenden Betriebskaufvorgänge."""

from datetime import date
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader
from testakte_file_filter import include_in_working_dump

ROOT = Path(__file__).resolve().parents[1]
CASES = {
    "insolvenzforderungen-handwerk-koeln": (60, 70, date(2026, 9, 10), "61_Betriebskaufvertrag_Entwurf_10-09-2026.docx", "62_Assetbestand_09-09-2026.xlsx"),
    "insolvenzforderungen-pistazienbrezeln-muenchen": (23, 32, date(2026, 8, 28), "24_Betriebskaufvertrag_Entwurf_28-08-2026.docx", "25_Assetbestand_28-08-2026.xlsx"),
}


def doc(path):
    with ZipFile(path) as z:
        tree = ET.fromstring(z.read("word/document.xml"))
    return " ".join(n.text or "" for n in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))


class AssetDealCases(unittest.TestCase):
    def test_native_inventory_and_readme_complete(self):
        for slug, (first, last, *_rest) in CASES.items():
            case = ROOT / "testakten" / slug
            files = [p for p in case.iterdir() if include_in_working_dump(p, case)]
            self.assertEqual(len(files), last, slug)
            self.assertEqual(sorted(int(p.name[:2]) for p in files), list(range(1, last + 1)))
            readme = (case / "README.md").read_text()
            for p in files:
                self.assertIn(p.name, readme)
                self.assertNotEqual(p.suffix, ".md")
                if int(p.name[:2]) >= first and p.suffix == ".docx":
                    text = doc(p)
                    self.assertGreater(len(text), 900, p.name)
                    self.assertNotRegex(text, r"(?i)Musterlösung|Lösungsmatrix|Formathinweis|Platzhalter|TODO")
                    self.assertNotIn("§", text)

    def test_drafts_are_unexecuted_and_have_editable_schedules(self):
        for slug, (_, _, _, contract, _) in CASES.items():
            text = doc(ROOT / "testakten" / slug / contract)
            self.assertGreater(len(text), 7500)
            for term in ("Anlage 1", "[", "01.10.2026", "Paragraf 613a BGB", "Paragraf 164 InsO", "Paragraf 1 Absatz 1a UStG"):
                self.assertIn(term, text)
            self.assertRegex(text, r"Nicht unterzeichnet|Keine Unterschrift")
            self.assertRegex(text, r"Zustimmung des jeweiligen Vertragspartners|Vertragspartner und beide Vertragsparteien")
            self.assertIn("zwingende", text.lower())

    def test_asset_workbooks_are_separate_raw_records(self):
        for slug, (_, _, _, _, name) in CASES.items():
            workbook = load_workbook(ROOT / "testakten" / slug / name, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Inventar", "Waren", "Kunden", "Verträge", "Kennzeichen"])
            for sheet in workbook:
                self.assertGreaterEqual(sheet.max_row, 8)
                self.assertEqual(sheet.max_column, 6)
                self.assertEqual(sheet.page_setup.orientation, "landscape")
                self.assertEqual(sheet.page_setup.fitToWidth, 1)
                ids = [row[0] for row in sheet.iter_rows(min_row=6, values_only=True)]
                self.assertEqual(len(ids), len(set(ids)))
                text = " ".join(str(c.value or "") for row in sheet for c in row)
                self.assertNotRegex(text, r"(?i)Prüfergebnis|anerkennen|bestreiten|Lösungsmatrix")
            for row in workbook["Waren"].iter_rows(min_row=6):
                self.assertIsInstance(row[2].value, (int, float))
                self.assertGreater(row[2].value, 0)
            workbook.close()

    def test_cologne_stock_reconciles_to_returns_and_usage(self):
        case = ROOT / "testakten/insolvenzforderungen-handwerk-koeln"
        wb = load_workbook(case / CASES[case.name][-1], data_only=True)
        quantities = {row[0]: row[2] for row in wb["Waren"].iter_rows(min_row=6, values_only=True)}
        self.assertEqual(quantities, {"W-01": 3000 - 2520, "W-02": 120 - 78, "W-03": 200 - 20 - 162, "W-04": 120 - 5 - 91, "W-05": 4})
        credit = doc(case / "39_Gutschrift_G-260422.docx")
        self.assertIn("20 Brandschutzplatten", credit)
        self.assertIn("5 Anschlussprofile", credit)
        self.assertIn("HS42-260615-07", doc(case / "61_Betriebskaufvertrag_Entwurf_10-09-2026.docx"))
        self.assertIn("20000,00", doc(case / "61_Betriebskaufvertrag_Entwurf_10-09-2026.docx"))
        wb.close()

    def test_munich_oven_and_new_pistachio_lot_remain_distinct(self):
        case = ROOT / "testakten/insolvenzforderungen-pistazienbrezeln-muenchen"
        wb = load_workbook(case / CASES[case.name][-1], data_only=True)
        serials = [r[2] for r in wb["Inventar"].iter_rows(min_row=6, values_only=True)]
        self.assertIn("SB-E3-240204", serials)
        self.assertNotIn("24118", serials)
        stock = list(wb["Waren"].iter_rows(min_row=6, values_only=True))[2]
        self.assertEqual((stock[2], stock[4]), (12, "KP-260820"))
        text = doc(case / "26_Bestandsaufnahme_Backstube_28-08-2026.docx")
        for token in ("KP-260602", "09.07.2026", "428,00", "WK-U4", "Backhof 14"):
            self.assertIn(token, text)
        wb.close()

    def test_correspondence_stays_within_case_cutoffs(self):
        for slug, (first, _, cutoff, *_rest) in CASES.items():
            for path in (ROOT / "testakten" / slug).glob("*.eml"):
                if int(path.name[:2]) < first:
                    continue
                msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                for field in ("Date", "From", "To", "Reply-To", "Message-ID", "Subject", "MIME-Version"):
                    self.assertTrue(msg[field], (path.name, field))
                self.assertLessEqual(parsedate_to_datetime(msg["Date"]).date(), cutoff)
                body = msg.get_body(preferencelist=("plain",)).get_content()
                self.assertGreater(len(body), 500)
                self.assertRegex(body, "[äöüß]")

    def test_combined_pdfs_include_every_new_source(self):
        for slug, (first, *_rest) in CASES.items():
            case = ROOT / "testakten" / slug
            pdf = PdfReader(case / "gesamt-pdf" / f"{slug}_gesamt.pdf")
            text = "\n".join(p.extract_text() or "" for p in pdf.pages)
            self.assertNotIn("Diese Testakte wurde", text)
            for p in case.iterdir():
                if p.name[:2].isdigit() and int(p.name[:2]) >= first:
                    self.assertIn(p.name, text)


if __name__ == "__main__":
    unittest.main()
