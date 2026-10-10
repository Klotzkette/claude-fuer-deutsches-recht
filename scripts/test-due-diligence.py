#!/usr/bin/env python3
"""Fachspezifische Quellen-, Spiegelungs- und Paketprüfungen für Due Diligence."""
from __future__ import annotations

import argparse
from collections import Counter
from decimal import Decimal
from openpyxl import load_workbook
from release_routing import case_asset_url, plugin_asset_url
from email import policy
from email.parser import BytesParser
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import sys
import unittest
import zipfile

import pdfplumber
from pypdf import PdfReader
import yaml
from dd_bankakte_pdf import assert_preserved, display, render as render_bank_workbook, source_cells
from testakte_disclaimer import NOTICE_BYTES, NOTICE_DE, NOTICE_EN
from testakte_file_filter import include_in_working_dump
from testakte_einzelpdf_common import document_arcname_pairs
from testakte_zip_common import working_dump_archive_pairs

ROOT = Path(__file__).resolve().parents[1]
NAME = 'due-diligence'
PLUGIN = ROOT / NAME
BANK = ROOT / 'testakten/dd-bankfiliale-verbraucherdarlehen-erfurt'
CASES = ('dd-arbeitsvertraege-innovation-berlin', 'dd-corporate-silberfalke', BANK.name)
DIST = None


def load_builder():
    spec = importlib.util.spec_from_file_location('dd_builder', ROOT / 'scripts/build-due-diligence-release.py')
    obj = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(obj)
    return obj


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class DueDiligenceTests(unittest.TestCase):
    def assert_bank_pdf_legible_and_complete(self, data):
        path = BANK / '11_Portfolio_und_Kaufpreis.xlsx'
        workbook, cells = source_cells(path)
        self.assertEqual(len(cells), 14633)
        assert_preserved(path, data)
        text = '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(data)).pages)
        groups = (
            ('2.1 Bestand:', '2.2 Bestand:', 'Portfolio', 'ABCDEFG', 1005),
            ('2.2 Bestand:', '3.1 Buchungen:', 'Portfolio', 'AHIJK', 1005),
            ('3.1 Buchungen:', '3.2 Buchungen:', 'Buchungen', 'ABCDE', 245),
            ('3.2 Buchungen:', '3.3 Buchungen:', 'Buchungen', 'ABFGHI', 245),
            ('3.3 Buchungen:', '4. Auswahl', 'Buchungen', 'ABJKLMN', 245),
        )
        for start, end, name, columns, last in groups:
            portion = text.split(start, 1)[1].split(end, 1)[0]
            compact = re.sub(r'\s+', '', portion)
            sheet = workbook[name]
            expected_amounts = Counter()
            for row in range(6, last + 1):
                values = [display(sheet[f'{column}{row}']) for column in columns]
                self.assertIn(re.sub(r'\s+', '', ''.join(values)), compact, f'{start} {row}')
                expected_amounts.update(value for value in values if re.fullmatch(r'-?[\d.]+,\d{2}', value))
            printed_amounts = Counter(re.findall(r'(?<![\w.,])-?[\d.]+,\d{2}(?![\d%])', portion))
            self.assertFalse(expected_amounts - printed_amounts, start)
            if name == 'Portfolio':
                self.assertEqual({f'APB-{i:04}' for i in range(1, 1001)}, set(re.findall(r'APB-\d{4}', portion)))
        with pdfplumber.open(io.BytesIO(data)) as document:
            for index, page in enumerate(document.pages, 1):
                self.assertAlmostEqual(page.width, 841.8898, places=2)
                self.assertAlmostEqual(page.height, 595.2756, places=2)
                self.assertTrue(page.chars, f'Leerseite {index}')
                for char in page.chars:
                    self.assertGreaterEqual(char['size'], 9, f'Schrift zu klein auf Seite {index}')
                    self.assertGreaterEqual(char['x0'], 35, f'Linker Rand {index}')
                    self.assertLessEqual(char['x1'], page.width - 35, f'Rechter Rand {index}')
                    self.assertGreaterEqual(char['top'], 25, f'Oberer Rand {index}')
                    self.assertLessEqual(char['bottom'], page.height - 14, f'Unterer Rand {index}')

    def test_bank_workbook_print_view_all_cells_rows_amounts_and_readable_geometry(self):
        path = BANK / '11_Portfolio_und_Kaufpreis.xlsx'
        before = sha(path)
        self.assert_bank_pdf_legible_and_complete(render_bank_workbook(path))
        self.assertEqual(before, sha(path), 'Die lesbare Druckansicht darf die Original-XLSX nicht ändern.')

    def test_manifests_eleven_skills_and_frontmatter(self):
        manifests = [json.loads((PLUGIN / p).read_text()) for p in
                     ('plugin.json', '.claude-plugin/plugin.json', '.codex-plugin/plugin.json')]
        self.assertEqual({(m['name'], m['version'], m['description']) for m in manifests},
                         {(NAME, '1.0.1', manifests[0]['description'])})
        skills = sorted((PLUGIN / 'skills').glob('*/SKILL.md'))
        self.assertEqual(len(skills), 11)
        for skill in skills:
            text = skill.read_text()
            frontmatter = yaml.safe_load(text.split('---', 2)[1])
            self.assertEqual(set(frontmatter), {'name', 'description'}, str(skill))
            self.assertLessEqual(len(frontmatter['description']), 360, str(skill))
            self.assertNotIn('§', frontmatter['description'])
            for n in range(1, 7):
                self.assertRegex(text, rf'(?m)^#{{1,3}} {n}[. ]', str(skill))
            for required in ('Times New Roman', '11 pt', 'ausformuliert'):
                self.assertIn(required, text, str(skill))
            for target in re.findall(r'\]\(([^ )]+)\)', text):
                if not target.startswith(('http:', 'https:', '#')):
                    self.assertTrue((skill.parent / target.split('#')[0]).exists(), f'{skill}: {target}')

    def test_commands_have_manufacturer_metadata_and_accept_arguments(self):
        names = {'dd', 'dd-schnellpruefung', 'dd-rueckfragen', 'dd-entscheidung'}
        commands = sorted((PLUGIN / 'commands').glob('*.md'))
        self.assertEqual({path.stem for path in commands}, names)
        for path in commands:
            text = path.read_text()
            self.assertTrue(text.startswith('---\n'), str(path))
            metadata = yaml.safe_load(text.split('---', 2)[1])
            self.assertIsInstance(metadata, dict, str(path))
            self.assertIsInstance(metadata.get('description'), str, str(path))
            self.assertTrue(metadata['description'].strip(), str(path))
            self.assertIsInstance(metadata.get('argument-hint'), str, str(path))
            self.assertIn('$ARGUMENTS', text.split('---', 2)[2], str(path))

    def test_plugin_links_are_self_contained(self):
        for path in PLUGIN.rglob('*.md'):
            for target in re.findall(r'\]\(([^ )]+)\)', path.read_text()):
                if target.startswith(('http:', 'https:', '#')):
                    continue
                resolved = (path.parent / target.split('#')[0]).resolve()
                self.assertTrue(resolved.is_relative_to(PLUGIN.resolve()), f'{path}: {target}')
                self.assertTrue(resolved.exists(), f'{path}: {target}')

    def test_three_prompt_pairs(self):
        for kind in ('werkstatt', 'schnellstart', 'hauptproblem'):
            source = PLUGIN / f'{NAME}-{kind}.md'
            data = source.read_bytes()
            self.assertEqual(data, source.with_suffix('.txt').read_bytes(), kind)
            if kind != 'werkstatt':
                self.assertLessEqual(len(data), 7500, kind)
            self.assertTrue(data.decode('utf-8').strip(), kind)

    def test_mirrors_have_exact_original_bytes_and_filter_selection(self):
        manifest = json.loads((ROOT / 'quality/due-diligence/spiegelung.json').read_text())
        self.assertEqual(len(manifest['cases']), 2)
        for case in manifest['cases']:
            original = ROOT / 'testakten' / case['original']
            mirror = ROOT / 'testakten' / case['mirror']
            tracked = {row['mirror'] for row in case['files']}
            native = {str(p.relative_to(ROOT)) for p in mirror.rglob('*') if p.is_file()
                      and p.name not in {'README.md', 'rubric.yaml'} and 'gesamt-pdf' not in p.parts}
            self.assertEqual(tracked, native)
            self.assertEqual(len(tracked), case['source_native_files'])
            exported = 0
            for row in case['files']:
                source, target = ROOT / row['original'], ROOT / row['mirror']
                self.assertEqual(sha(source), row['sha256'], str(source))
                self.assertEqual(sha(target), row['sha256'], str(target))
                self.assertEqual(target.stat().st_size, row['bytes'])
                self.assertEqual(include_in_working_dump(source, original), row['working_export'])
                self.assertEqual(include_in_working_dump(target, mirror), row['working_export'])
                exported += row['working_export']
            self.assertEqual(exported, case['working_files'])

    def test_fifty_employment_contracts_and_two_director_agreements(self):
        directory = ROOT / 'testakten' / CASES[0]
        contracts = sorted(directory.glob('*_arbeitsvertrag_*.docx'))
        self.assertEqual(len(contracts), 50)
        self.assertEqual([int(p.name.split('_', 1)[0]) for p in contracts], list(range(1, 51)))
        self.assertEqual(len(list(directory.glob('*_dienstvertrag_*.docx'))), 2)
        self.assertTrue(list(directory.glob('*.xlsx')))
        self.assertGreaterEqual(len(list(directory.glob('*.eml'))), 10)

    def test_silberfalke_is_declared_as_review_of_historical_analysis(self):
        readme = (ROOT / 'testakten' / CASES[1] / 'README.md').read_text()
        for phrase in ('Lernakte', 'kein blinder Leistungstest', 'keine verifizierte Rechtsquelle',
                       'statuskarte_erwartung.docx', 'dd_report_cover_executive_summary.docx'):
            self.assertIn(phrase, readme)

    def test_case_readmes_disclaimer_and_distinct_representations(self):
        for case in CASES:
            readme = (ROOT / 'testakten' / case / 'README.md').read_text()
            for line in (NOTICE_DE, NOTICE_EN):
                if line.strip():
                    self.assertIn(line.strip(), readme, case)
            self.assertIn('due-diligence-v1.0.1', readme, case)
            for suffix in ('.zip', '-einzelpdfs.zip', '_gesamt.pdf'):
                self.assertIn(case + suffix, readme, case)

    def test_bank_source_formats_and_mime_integrity(self):
        files = [p for p in BANK.rglob('*') if include_in_working_dump(p, BANK)]
        self.assertGreaterEqual(len(files), 60, 'Die 20 Darlehen brauchen echte Unterlagen und Korrespondenz.')
        for suffix in ('.docx', '.pdf', '.eml', '.xlsx'):
            self.assertTrue(any(p.suffix == suffix for p in files), suffix)
        mails = [p for p in files if p.suffix == '.eml']
        self.assertGreaterEqual(len(mails), 10)
        for path in mails:
            mail = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            for header in ('From', 'To', 'Subject', 'Date', 'Message-ID', 'MIME-Version'):
                self.assertTrue(mail[header], f'{path}: {header}')
            self.assertFalse(mail.defects, str(path))
            for part in mail.iter_attachments():
                self.assertTrue(part.get_payload(decode=True), str(path))
        self.assertFalse(any(p.suffix == '.md' for p in files))

    def test_bank_portfolio_and_twenty_lived_accounts_reconcile_in_cents(self):
        data = json.loads((ROOT / 'scripts/data/dd-bankakte.json').read_text())
        portfolio = data['portfolio']
        self.assertEqual(len(portfolio), 1000)
        self.assertEqual({row['id'] for row in portfolio}, {f'APB-{n:04d}' for n in range(1, 1001)})
        self.assertEqual(len(data['cases']), 20)
        self.assertEqual({row['id'] for row in data['cases']}, {f'APB-{n:04d}' for n in range(1, 21)})
        for row in portfolio:
            for key in ('principal', 'interest', 'costs', 'balance', 'price'):
                self.assertIs(type(row[key]), int, f"{row['id']}: {key} muss Ganzzahlcent sein")
            self.assertEqual(row['balance'], row['principal'] + row['interest'] + row['costs'])
            self.assertEqual(row['price'], round(row['principal'] * row['seller_price_ratio']))
        for key, total in data['totals'].items():
            self.assertEqual(total, sum(row[key] for row in portfolio), key)
        mapping = {row['id']: row for row in portfolio}
        for case in data['cases']:
            self.assertEqual(len(case['journal']), 12, case['id'])
            previous_principal = case['opening_principal']
            previous_interest = case['opening_interest']
            previous_costs = 0
            for entry in case['journal']:
                self.assertEqual(entry['principal_start'], previous_principal)
                self.assertEqual(entry['principal_end'], previous_principal - entry['paid_principal'])
                self.assertEqual(entry['interest_end'], previous_interest + entry['interest_charge'] - entry['paid_interest'])
                self.assertEqual(entry['costs_end'], previous_costs + entry['costs_charge'] - entry['paid_costs'])
                self.assertEqual(entry['cash'], sum(entry[key] for key in ('paid_principal', 'paid_interest', 'paid_costs')))
                self.assertEqual(entry['total_end'], sum(entry[key] for key in ('principal_end', 'interest_end', 'costs_end')))
                previous_principal, previous_interest, previous_costs = (entry[key] for key in ('principal_end', 'interest_end', 'costs_end'))
            for key in ('principal', 'interest', 'costs', 'balance'):
                self.assertEqual(case[key], mapping[case['id']][key], case['id'])
            self.assertEqual(previous_principal, case['principal'])
            self.assertEqual(previous_interest, case['interest'])
            self.assertEqual(previous_costs, case['costs'])
            for suffix in ('01_Darlehensvertrag.docx', '02_Kontojournal.pdf', '03_Schreiben.pdf', '04_Korrespondenz.eml'):
                self.assertTrue((BANK / f"{case['id']}_{suffix}").is_file(), f"{case['id']}_{suffix}")
        self.assertEqual(len(list(BANK.glob('APB-*_01_Darlehensvertrag.docx'))), 20)
        self.assertEqual(data['general_ledger_net'], data['totals']['balance'] + sum(row['amount'] for row in data['adjustments']))
        self.assertEqual(data['adjustments'][0]['loan'], 'APB-0017')
        self.assertEqual(data['adjustments'][0]['amount'], -72000)

    def test_bank_workbook_cached_values_preserve_raw_totals_and_open_review(self):
        data = json.loads((ROOT / 'scripts/data/dd-bankakte.json').read_text())
        path = BANK / '11_Portfolio_und_Kaufpreis.xlsx'
        workbook = load_workbook(path, data_only=True)
        sheet = workbook['Ueberblick']
        self.assertEqual([sheet[c].value for c in ('B7', 'B8', 'B9')], [1000, 20, 980])
        for cell, key in (('B10', 'principal'), ('B11', 'interest'), ('B12', 'costs'), ('B13', 'balance'), ('E7', 'price')):
            self.assertEqual(Decimal(str(sheet[cell].value)).quantize(Decimal('0.01')) * 100, data['totals'][key], cell)
        self.assertEqual(Decimal(str(sheet['B18'].value)) * 100, -72000)
        self.assertEqual(Decimal(str(sheet['B20'].value)) * 100, data['general_ledger_net'])
        self.assertEqual(sheet['B24'].value, 0)
        self.assertEqual(sheet['B26'].value, 20, 'Keine verdeckt vorgegebene abgeschlossene Prüfung.')
        for tab in workbook:
            for row in tab:
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'e', f'{tab.title}!{cell.coordinate}')
        formulas = load_workbook(path, data_only=False)
        for row in range(6, 1006):
            self.assertEqual(formulas['Portfolio'][f'G{row}'].data_type, 'f')
            self.assertEqual(formulas['Portfolio'][f'I{row}'].data_type, 'f')
        for row in range(6, 246):
            self.assertEqual(formulas['Buchungen'][f'I{row}'].data_type, 'f')
            self.assertEqual(formulas['Buchungen'][f'J{row}'].data_type, 'f')
            self.assertEqual(formulas['Buchungen'][f'M{row}'].data_type, 'f')

    def test_component_routes_do_not_displace_existing_case_downloads(self):
        self.assertIn('/due-diligence-v1.0.1/', plugin_asset_url(NAME))
        for case in CASES:
            for suffix in ('', '-einzelpdfs'):
                self.assertIn('/due-diligence-v1.0.1/', case_asset_url(case, suffix))
        for case in ('weg-lindenhof-jahresabrechnung-2025', 'weg-sonnenwinkel-bettwanzen', 'weg-spreebogen-mieterumlage-belege'):
            for suffix in ('', '-einzelpdfs'):
                self.assertIn('/akten-v445.35.3/', case_asset_url(case, suffix))

    def test_installable_archive_sources_do_not_include_test_cases_or_prompts(self):
        paths = load_builder().plugin_sources()
        self.assertTrue(paths)
        for p in paths:
            self.assertNotIn('testakten', p.parts)
            self.assertNotIn(p.suffix.lower(), {'.eml', '.pdf', '.docx'})
            self.assertFalse(re.search(r'-(werkstatt|schnellstart|hauptproblem)\.(md|txt)$', p.name))

    def test_release_assets_checksums_and_exact_plugin_payload(self):
        if DIST is None:
            self.skipTest('Paketprüfung erst nach explizitem --dist.')
        builder = load_builder()
        expected = builder.asset_names() | {'checksums-sha256.txt'}
        self.assertEqual({p.name for p in DIST.iterdir()}, expected)
        self.assertEqual(len(expected), 21)
        checksum_rows = (DIST / 'checksums-sha256.txt').read_text().splitlines()
        checksums = dict(line.split('  ', 1)[::-1] for line in checksum_rows)
        self.assertEqual(set(checksums), expected - {'checksums-sha256.txt'})
        for name, digest in checksums.items():
            self.assertEqual(sha(DIST / name), digest, name)
        for portable in (False, True):
            prefix = f'{NAME}/' if portable else ''
            with zipfile.ZipFile(DIST / f'{NAME}{"-portable" if portable else ""}.zip') as archive:
                self.assertIsNone(archive.testzip())
                paths = {p.relative_to(PLUGIN).as_posix(): p for p in builder.plugin_sources()}
                paths['LICENSE'] = ROOT / 'LICENSE'
                self.assertEqual(set(archive.namelist()), {prefix + p for p in paths})
                for name, path in paths.items():
                    self.assertEqual(archive.read(prefix + name), path.read_bytes())

    def test_three_cases_complete_same_pages_and_no_warning_inside_pdfs(self):
        if DIST is None:
            self.skipTest('Paketprüfung erst nach explizitem --dist.')
        for case in CASES:
            directory = ROOT / 'testakten' / case
            with zipfile.ZipFile(DIST / f'testakte-{case}.zip') as archive:
                self.assertIsNone(archive.testzip())
                pairs = working_dump_archive_pairs(directory, include_gesamt_pdf=True)
                self.assertEqual(set(archive.namelist()), {'README.txt', *(name for _, name in pairs)})
                self.assertTrue(archive.read('README.txt').startswith(NOTICE_BYTES))
                self.assertFalse(any('/' in name or name.endswith('.md') for name in archive.namelist()))
                for path, name in pairs:
                    self.assertEqual(archive.read(name), path.read_bytes(), name)
            with zipfile.ZipFile(DIST / f'testakte-{case}-einzelpdfs.zip') as archive:
                pairs = document_arcname_pairs(directory)
                self.assertEqual(set(archive.namelist()), {'README.txt', *(name for _, name in pairs)})
                self.assertTrue(archive.read('README.txt').startswith(NOTICE_BYTES))
                pages = []
                for _, name in pairs:
                    reader = PdfReader(io.BytesIO(archive.read(name)))
                    self.assertGreater(len(reader.pages), 0, name)
                    pages.extend(p.extract_text() or '' for p in reader.pages)
                combined = PdfReader(DIST / f'{case}_gesamt.pdf')
                self.assertEqual(len(pages), len(combined.pages), case)
                for expected, page in zip(pages, combined.pages):
                    actual = page.extract_text() or ''
                    self.assertEqual(expected, actual)
                    self.assertNotIn('Diese Testakte wurde mit KI generiert', actual)
                    self.assertNotIn('This test case file was generated with AI', actual)

    def test_prompt_reading_versions_do_not_lose_text(self):
        if DIST is None:
            self.skipTest('Paketprüfung erst nach explizitem --dist.')
        for kind in ('werkstatt', 'schnellstart', 'hauptproblem'):
            stem = f'{NAME}-{kind}'
            original = (PLUGIN / f'{stem}.md').read_bytes()
            for ext in ('md', 'txt'):
                self.assertEqual((DIST / f'{stem}.{ext}').read_bytes(), original)
            text = '\n'.join(page.extract_text() or '' for page in PdfReader(DIST / f'{stem}.pdf').pages)
            # Lange amtliche URLs dürfen im PDF umbrechen. Vollständigkeit
            # zuerst zeichengetreu ohne Umbruch-Leerraum prüfen; danach die
            # übrigen Wörter zählen, statt URL-Silben als Textverlust zu melden.
            original_text = original.decode()
            comparison_text = text
            for url in sorted(set(re.findall(r'https?://[^\s)]+', original_text)), key=len, reverse=True):
                url = url.rstrip('.,;')
                wrapped = r'\s*'.join(re.escape(character) for character in url)
                self.assertRegex(comparison_text, wrapped, f'{kind}: amtliche URL unvollständig: {url}')
                comparison_text = re.sub(wrapped, ' ', comparison_text)
                original_text = original_text.replace(url, ' ')
            words = lambda data: Counter(re.findall(r'[\w§]+', data.casefold()))
            missing = words(original_text) - words(comparison_text)
            self.assertFalse(missing, f'{kind}: {missing}')
            self.assertNotIn('\ufffd', text)
            self.assertNotIn('\u25a0', text)

    def test_published_bank_workbook_uses_readable_print_view(self):
        if DIST is None:
            self.skipTest('Paketprüfung erst nach explizitem --dist.')
        with zipfile.ZipFile(DIST / f'testakte-{BANK.name}-einzelpdfs.zip') as archive:
            pairs = dict((path.name, arcname) for path, arcname in document_arcname_pairs(BANK))
            data = archive.read(pairs['11_Portfolio_und_Kaufpreis.xlsx'])
        self.assert_bank_pdf_legible_and_complete(data)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dist', type=Path)
    args, extra = parser.parse_known_args()
    DIST = args.dist
    unittest.main(argv=[sys.argv[0], *extra])
