#!/usr/bin/env python3
"""Rendert vier redaktionell verfasste Haftakten ohne Lösungsunterlagen."""

import csv
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader, PdfWriter

from testakte_disclaimer import NOTICE_BYTES, NOTICE_MARKDOWN
from testakte_office_pdf import render_office_batch

ROOT = Path(__file__).resolve().parents[1]
DATA = Path(__file__).with_name("data") / "untersuchungshaft-akten.json"
DOWNLOADS = ROOT / "untersuchungshaft/testakte/downloads"
TAG = "untersuchungshaft-v1.0.0"
BASE = f"https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/{TAG}"
STAMP = datetime(2026, 10, 16, 8)


def module(name):
    spec = importlib.util.spec_from_file_location(name.replace("-", "_"), Path(__file__).with_name(name + ".py"))
    loaded = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(loaded)
    return loaded


def archive_bytes(entries):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, data in sorted(entries.items()):
            if "/" in name or "\\" in name or not name.isascii():
                raise ValueError(f"Ungültiger flacher Dateiname: {name}")
            info = zipfile.ZipInfo(name, (2026, 10, 16, 8, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            archive.writestr(info, data)
    return out.getvalue()


def document(path, item):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.3)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(5 if item.get("compact") else 7)
        style.paragraph_format.line_spacing = 1.0 if item.get("compact") else 1.06
        style.paragraph_format.widow_control = True
        fonts = style.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if key.endswith("Theme"):
                del fonts.attrib[key]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + script), "Times New Roman")
    doc.styles["Title"].font.size = Pt(15)
    doc.styles["Title"].font.bold = True
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for value in (item["author"], item["to"], item["date"]):
        doc.add_paragraph(value)
    doc.add_paragraph(item["title"], "Title")
    for block in item["body"].split("\n\n"):
        heading = bool(re.match(r"^\d+(?:\.\d+)* ", block)) and "\n" not in block and len(block) < 90
        doc.add_paragraph(block, "Heading 1" if heading else "Normal")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Seite ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = doc.core_properties.last_modified_by = "Klotzkette"
    doc.core_properties.title = item["title"]
    doc.core_properties.language = "de-DE"
    doc.core_properties.created = doc.core_properties.modified = STAMP
    buffer = io.BytesIO()
    doc.save(buffer)
    # OOXML-Inhalt und ZIP-Zeitstempel bleiben bei gleichem Text reproduzierbar.
    with zipfile.ZipFile(buffer) as source:
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, (2026, 10, 16, 8, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                target.writestr(info, source.read(name))
        path.write_bytes(out.getvalue())


def write_original(path, item, slug):
    suffix = path.suffix
    if suffix == ".docx":
        document(path, item)
    elif suffix == ".eml":
        mail = EmailMessage(policy=SMTP)
        mail["From"], mail["To"] = item["author"], item["to"]
        mail["Date"] = format_datetime(datetime.fromisoformat(item["date"]))
        mail["Subject"] = item["title"]
        mail["Message-ID"] = f"<{slug}-{path.stem}@akten.example>"
        mail["Content-Language"] = "de-DE"
        mail.set_content(item["body"], charset="utf-8")
        for name in item.get("attachments", []):
            attachment = path.parent / name
            if attachment.name != name or attachment.suffix != ".docx":
                raise ValueError(f"Ungültige Mailanlage: {name}")
            mail.add_attachment(attachment.read_bytes(), maintype="application",
                                subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                                filename=name)
        if item.get("attachments"):
            mail.set_boundary(f"anlage-{slug}-{path.stem}")
        path.write_bytes(mail.as_bytes())
    elif suffix == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle, delimiter=";").writerows(item["rows"])
    elif suffix == ".txt":
        path.write_text("\n\n".join(item[k] for k in ("author", "to", "date", "title", "body")) + "\n", encoding="utf-8")
    else:
        raise ValueError(suffix)


def build():
    renderer = module("build-testakten-einzelpdf-zips")
    cases = json.loads(DATA.read_text(encoding="utf-8"))
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    report = []
    for case in cases:
        slug = case["slug"]
        directory = ROOT / "testakten" / slug
        directory.mkdir(parents=True, exist_ok=True)
        paths = []
        for item in case["files"]:
            path = directory / item["name"]
            if path.name != item["name"] or not path.name.isascii():
                raise ValueError(f"Ungültiger Originalname: {item['name']}")
            write_original(path, item, slug)
            paths.append(path)
        cache = render_office_batch(paths)
        missing = [p.name for p in paths if p.suffix == ".docx" and p not in cache]
        if missing:
            raise RuntimeError(f"Layoutgetreue Office-PDFs fehlen: {missing}")
        single = {"README.txt": NOTICE_BYTES}
        original = {"README.txt": NOTICE_BYTES}
        writer = PdfWriter()
        for path in paths:
            pdf = renderer.render_document_pdf(path, directory, cache)
            if not pdf:
                raise RuntimeError(f"Leeres PDF: {path}")
            single[path.stem + ".pdf"] = pdf
            original[path.name] = path.read_bytes()
            start = len(writer.pages)
            for page in PdfReader(io.BytesIO(pdf)).pages:
                writer.add_page(page)
            writer.add_outline_item(path.name, start)
        writer.add_metadata({"/Title": case["title"], "/Author": "Klotzkette"})
        total = directory / "gesamt-pdf" / f"{slug}_gesamt.pdf"
        total.parent.mkdir(exist_ok=True)
        with total.open("wb") as handle:
            writer.write(handle)
        original[total.name] = total.read_bytes()
        for suffix, entries in (("", original), ("-einzelpdfs", single)):
            (DOWNLOADS / f"testakte-{slug}{suffix}.zip").write_bytes(archive_bytes(entries))
        rows = "\n".join(f"| [{d['name']}]({d['name']}) | {d['title']} |" for d in case["files"])
        readme = f"""# 1. {case['title']}

{case['summary']}

## 1.1. Einstieg und Material

Stand {case['date']}. Neun eigenständige Unterlagen: fünf Word-Dokumente, zwei E-Mails, eine Tabelle und eine Textdatei. Beginnen Sie mit der ersten gerichtlichen oder polizeilichen Unterlage und wählen Sie Ihre Rolle. Die vorhandenen Schreiben sind Verfahrensmaterial, keine vorweggenommene Lösung. Eine Entscheidung über den nächsten Antrag ist nicht beigelegt.

## 1.2. Downloads

{NOTICE_MARKDOWN}

| Format | Download |
| --- | --- |
| Gesamt-PDF mit Lesezeichen | [Gesamtakte](gesamt-pdf/{slug}_gesamt.pdf) |
| Jede Unterlage als eigenes PDF | [Einzel-PDF-ZIP]({BASE}/testakte-{slug}-einzelpdfs.zip) |
| Word, E-Mails, Tabelle und Text | [Originalformat-ZIP]({BASE}/testakte-{slug}.zip) |

Beide ZIPs sind flach und enthalten den zweisprachigen Hinweis in README.txt. Das Originalformat-ZIP enthält zusätzlich das Gesamt-PDF, aber kein Markdown. Die PDF-Sammlung enthält genau ein PDF je Originalunterlage. Mehrseitige Unterlagen bleiben jeweils zusammen; verschiedene Originale werden nicht zu einem Einzel-PDF vermischt.

English: Nine separate source documents, no answer key. Choose the combined reading PDF, the flat archive of individual PDFs, or the flat archive of native documents. The native archive also includes the combined PDF. Notices accompany the downloads and sit in README.txt inside both archives, not within the PDFs.

## 1.3. Unterlagen

| Datei | Gegenstand |
| --- | --- |
{rows}

## 1.4. Einordnung

<!-- reserved-example-contacts -->

Personen, private Stellen, Adressen und Vorgangsnummern sind erfunden; die genannten Gerichte und Anstalten dienen allein als örtlicher Rahmen. Die Schriftstücke sind keine echten Behördenentscheidungen. Kontaktadressen mit reservierten Beispieldomains sind nicht für die Einreichung bestimmt. Die Akte enthält keine echten Gesundheits- oder Ermittlungsdaten.

[Untersuchungshaft](../../untersuchungshaft/README.md) · [Alle Testakten](../README.md) · [Repository](../../README.md) · [Download-Index](../../ASSET_INDEX.md)
"""
        (directory / "README.md").write_text(readme, encoding="utf-8")
        # Die Rubrik kontrolliert nur den Bestand, nicht die juristische Lösung.
        rubric = f"name: {case['title']}\nplugin: untersuchungshaft\nstand: '{case['date']}'\nchecks:\n"
        for index, path in enumerate(paths, 1):
            rubric += f"- id: dokument-{index}\n  check_type: file_exists\n  description: Eigenständiges Original vorhanden.\n  path: {path.name}\n"
        (directory / "rubric.yaml").write_text(rubric, encoding="utf-8")
        report.append({"slug": slug, "documents": len(paths), "pages": len(writer.pages),
                       "originals": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}})
        print(f"{slug}: {len(paths)} Originale, {len(writer.pages)} PDF-Seiten, zwei flache ZIPs", flush=True)
    out = ROOT / "quality/untersuchungshaft"
    out.mkdir(parents=True, exist_ok=True)
    (out / "aktenbestand.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build()
