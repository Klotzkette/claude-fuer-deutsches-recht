#!/usr/bin/env python3
"""Native Quellen des Hannoveraner Gesellschaftsmandats; keine Release-Eingriffe."""

import argparse
import csv
import json
import os
import re
import subprocess
import sys
import tempfile
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime, parseaddr
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw

from akten_build_runtime import node_binary, screen_font
from akten_docx_format import separate_section_headings
from testakte_office_pdf import render_office_batch

ROOT = Path(__file__).resolve().parents[1]
SLUG = 'wirtschaftsanwalt-gesellschafterkonflikt-handwerk-hannover'
OUT = ROOT / 'testakten' / SLUG
QA = Path(os.environ['HANNOVER_AKTEN_QA']) if os.environ.get('HANNOVER_AKTEN_QA') else Path(tempfile.mkdtemp(prefix='wirtschaftsanwalt-hannover-qa-'))
SOURCES = []
PDF_DOCS = []
WARN_DE = 'Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.'
WARN_EN = 'This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'

FIRMA = ('Niemeyer & Brandes Gebäudetechnik GmbH', 'Hägenstraße 38, 30559 Hannover',
         'Telefon 0511 763920-0 · buero@nb-gebaeudetechnik.de',
         'Sitz Hannover · Amtsgericht Hannover HRB 218764 · Geschäftsführer Lars Niemeyer und Felix Brandes')
KUNDE = ('Hofgarten Gewerbehöfe GmbH', 'Bemeroder Straße 72, 30559 Hannover',
         'Projektverwaltung · projekt@hofgarten-gewerbe.de · 0511 489730-12',
         'Sitz Hannover · Amtsgericht Hannover HRB 223681 · Geschäftsführerin Maren Rösler')
LIEFERANT = ('Rößler Haustechnik Großhandel GmbH', 'Dieselstraße 17, 30827 Garbsen',
             'Debitorenbuchhaltung · forderungen@roessler-haustechnik.de · 05131 709620-0',
             'Sitz Garbsen · Amtsgericht Hannover HRB 219583 · Geschäftsführer Volker Rößler')
BANK = ('Leine Gewerbebank eG', 'Firmenkunden Hannover · Theaterstraße 19, 30159 Hannover',
        'Carolin Mertens · c.mertens@leine-gewerbebank.de · 0511 817260-41',
        'Sitz Hannover · Amtsgericht Hannover GnR 2047 · Vorstand Judith Arndt und Tobias Henke')


def register(name):
    assert name not in SOURCES
    SOURCES.append(name)
    return OUT / name


def document(name, sender, recipient, title, date, reference, body, closing):
    """Office erzeugt auch die Brief-PDFs; Quellzwillinge bleiben im temporären QA-Pfad."""
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2.9), Cm(2.1)
    s.left_margin, s.right_margin = Cm(2.2), Cm(2.1)
    s.header_distance, s.footer_distance = Cm(.8), Cm(.8)
    for name_style in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
        style = d.styles[name_style]
        style.font.name, style.font.size = 'Times New Roman', Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if key.endswith('Theme'):
                del fonts.attrib[key]
        for script in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + script), 'Times New Roman')
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.03
    d.styles['Title'].font.size = Pt(16)
    d.styles['Title'].font.bold = True
    for h in ['Heading 1', 'Heading 2']:
        d.styles[h].font.bold = True
        d.styles[h].paragraph_format.space_before = Pt(8)
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    header = s.header.paragraphs[0]
    run = header.add_run(sender[0] + '\n')
    run.bold = True
    header.add_run(sender[1] + '\n' + sender[2])
    header.paragraph_format.space_after = Pt(0)
    footer = s.footer.paragraphs[0]
    footer.text = sender[3] + '\n' + reference + ' · Seite '
    footer.paragraph_format.space_after = Pt(0)
    for run in footer.runs:
        run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    d.add_paragraph(recipient)
    d.add_paragraph(date + '\nBezug: ' + reference)
    d.add_paragraph(title, style='Title')
    for block in body.strip().split('\n\n'):
        block = block.strip()
        if block == '---':
            d.add_page_break()
        elif re.match(r'^\d+(?:\.\d+)*\s', block) and '\n' not in block and len(block) < 105:
            d.add_paragraph(block, style='Heading 2' if '.' in block.split()[0] else 'Heading 1')
        else:
            p = d.add_paragraph(block)
            p.paragraph_format.widow_control = True
    p = d.add_paragraph(closing)
    p.paragraph_format.keep_together = True
    # Der letzte Sachabsatz bleibt bei Unterschrift bzw. Zeichnungsvermerk.
    if len(d.paragraphs) > 1:
        d.paragraphs[-2].paragraph_format.keep_with_next = True
    separate_section_headings(d)
    d.core_properties.author = sender[0]
    d.core_properties.title = title
    d.core_properties.language = 'de-DE'
    d.core_properties.created = datetime(2026, 9, 23, 14)
    d.core_properties.modified = datetime(2026, 9, 23, 14)
    target = register(name)
    if target.suffix == '.pdf':
        staged = QA / 'office' / target.with_suffix('.docx').name
        d.save(staged)
        PDF_DOCS.append((staged, target))
    else:
        d.save(target)


def mail(name, sender, recipient, moment, subject, body, cc=None, reply=None):
    msg = EmailMessage(policy=SMTP)
    address = parseaddr(sender)[1]
    domain = address.split('@')[1]
    stamp = datetime.fromisoformat(moment + '+02:00')
    msg['Return-Path'] = '<' + address + '>'
    destination_domain = parseaddr(recipient)[1].split('@')[1]
    msg['Received'] = f'from mail.{domain} by mx.{destination_domain} with ESMTPS id {name[:2]}202609; {format_datetime(stamp)}'
    msg['From'], msg['To'] = sender, recipient
    if cc:
        msg['Cc'] = cc
    msg['Date'] = format_datetime(stamp)
    msg['Subject'] = subject
    msg['Message-ID'] = f'<{name[:-4]}@{domain}>'
    if reply:
        msg['In-Reply-To'] = reply
        msg['References'] = reply
    msg['X-Mailer'] = 'Microsoft Outlook 16.0'
    msg['Content-Language'] = 'de-DE'
    msg.set_content(body.strip() + '\n', charset='utf-8')
    register(name).write_bytes(msg.as_bytes())


def plain(name, body):
    register(name).write_text(body.strip() + '\n', encoding='utf-8')


def csvfile(name, labels, rows):
    with register(name).open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle, delimiter=';', quoting=csv.QUOTE_ALL)
        writer.writerow(labels)
        writer.writerows(rows)


def corporate_sources():
    document('01_Mandatsanfrage_2026-09-23.docx', FIRMA,
             'Rechtsanwältin Dr. Anna Seifert\nKanzlei Seifert und Voß\nAn der Börse 7, 30159 Hannover',
             'Mandat für unsere Gesellschaft', 'Hannover, 23. September 2026', 'NB / Gesellschaft und Projekt HG-26-041', '''Sehr geehrte Frau Dr. Seifert,

bitte übernehmen Sie die Beratung der Niemeyer & Brandes Gebäudetechnik GmbH. Felix und ich tragen diesen Auftrag gemeinsam. Vertragspartnerin Ihrer Kanzlei soll nur die Gesellschaft sein. Unsere privaten Vorstellungen zum Anteilsverkauf sind in den beigefügten Nachrichten enthalten; wir wissen noch nicht, wer uns dazu jeweils persönlich beraten wird.

Wir beschäftigen 15 Arbeitnehmerinnen und Arbeitnehmer: elf Monteure, zwei Auszubildende, Bettina Kruse im kaufmännischen Bereich und Eva Schünemann im Büro. Hinzu kommen wir beide als Geschäftsführer. Der Betrieb läuft weiter. Am Gewerbehof in Bemerode sind noch Inbetriebnahme und Einweisung zu erledigen. Der Kunde hat von der zweiten Abschlagsrechnung 20.000,00 EUR nicht bezahlt und rügt Geräusche sowie Feuchtigkeit. Die Schlussrate ist noch nicht abgerechnet. Frau Kruse hat die bisherige Zahlung und alle offenen Posten getrennt zusammengestellt.

Rößler fordert seit Anfang September 29.274,00 EUR. Morgen läuft die letzte im Mahnschreiben genannte Frist ab. Für Ende September stehen Sozialbeiträge und Löhne an. Die Bank hat mehr Kredit in Aussicht gestellt, aber weitere Bürgschaften und Unterlagen verlangt. Felix möchte seine persönliche Haftung gerade nicht erhöhen. Die Excel-Datei beruht auf dem Bankstand von heute Morgen; Kundenprognosen sind noch keine Gutschriften.

Bitte klären Sie, wie wir für die Gesellschaft jetzt gegenüber Kunde, Lieferant und Bank auftreten können, welche zusätzlichen Zahlungsdaten Sie kurzfristig benötigen und wer die jeweiligen Erklärungen abgeben sollte. Zum Kunden brauchen wir eine abgestimmte Reaktion vor dem Ortstermin am Freitag, 25. September, 8:30 Uhr. Eine Zahlung oder ein Anerkenntnis haben wir Ihrer Kanzlei nicht zur Ausführung übertragen.

Außerdem möchten wir wissen, wie eine geordnete Trennung vorbereitet werden kann, ohne die Gesellschaft zur Käuferin von Felix' Anteil zu erklären. Bitte prüfen Sie unseren noch nicht versandten Einladungsentwurf für den 9. Oktober. Unter anderem sollen Finanzierung, Geschäftsführungsaufgaben und die verschiedenen Gehaltsvorschläge besprochen werden. Eine Abberufung, ein Verkauf oder eine Gehaltsänderung ist bisher nicht beschlossen. Bitte teilen Sie uns mit, welche Unterlagen für diese Punkte noch fehlen und wie Sie das Gesellschaftsmandat von unseren Einzelinteressen abgrenzen.

Frau Kruse darf Ihnen die Buchhaltung erläutern. Für Rückfragen sind wir heute ab 16:00 Uhr gemeinsam erreichbar. Bitte bestätigen Sie vor Beginn Ihr Honorar und den Umfang der übernommenen Angelegenheiten. Die elektronische Akte übermitteln wir heute mit 26 Einzeldateien.''',
             'Mit freundlichen Grüßen\nLars Niemeyer                         Felix Brandes\nGeschäftsführer                         Geschäftsführer')
    document('02_Register_und_Anteilsstand_2026-09-22.pdf', FIRMA,
             'An die Geschäftsführung\nGesellschaftsordner / Registerunterlagen',
             'Abschrift aus dem Registerordner', 'Hannover, 22. September 2026', 'HRB 218764 / Ordner Register, Blätter 1 und 6', '''Sehr geehrter Herr Niemeyer, sehr geehrter Herr Brandes,

ich habe die nachstehenden Angaben aus dem am 22. September 2026 um 09:12 Uhr im Büro abgelegten aktuellen Registerausdruck und aus der dort gespeicherten Gesellschafterliste vom 18. Mai 2018 übertragen. Diese Büroabschrift ersetzt keinen beglaubigten Registerausdruck. Die beiden Herkunftsunterlagen liegen im Gesellschaftsordner, Registerfach 1.

1 Firma und Gegenstand

Registergericht ist das Amtsgericht Hannover, Registerblatt HRB 218764. Die Firma lautet Niemeyer & Brandes Gebäudetechnik GmbH, Sitz Hannover, Geschäftsanschrift Hägenstraße 38, 30559 Hannover. Gegenstand sind Installation, Wartung und Instandsetzung von Heizungs-, Sanitär- und Lüftungsanlagen sowie der Handel mit zugehörigen Bauteilen. Das Stammkapital beträgt 50.000,00 EUR. Gesellschaftsvertrag vom 18. Mai 2018; erste Eintragung am 7. Juni 2018.

2 Vertretung und Personen

Allgemeine Vertretungsregelung laut Ausdruck: Ist nur ein Geschäftsführer bestellt, vertritt er die Gesellschaft allein. Sind mehrere Geschäftsführer bestellt, wird die Gesellschaft durch zwei Geschäftsführer oder durch einen Geschäftsführer gemeinsam mit einem Prokuristen vertreten. Einzelvertretungsbefugnis kann erteilt werden.

Als Geschäftsführer sind Lars Niemeyer, Hannover, und Felix Brandes, Hannover, eingetragen. Bei beiden ist ausdrücklich Einzelvertretungsbefugnis vermerkt. Eine Befreiung von den Beschränkungen des Paragrafen 181 BGB ist bei keinem der beiden eingetragen. Bettina Kruse, Hannover, ist seit 10. Februar 2023 mit Einzelprokura eingetragen. Eine besondere Grundstücksbefugnis ist nicht vermerkt. Weitere Prokuristen und Niederlassungen sind nicht aufgeführt.

3 Gesellschafterliste

Die gesonderte Gesellschafterliste weist Geschäftsanteil Nr. 1 mit 25.000,00 EUR für Lars Niemeyer und Geschäftsanteil Nr. 2 mit 25.000,00 EUR für Felix Brandes aus. Die Beteiligung beträgt jeweils 50 Prozent; die Summe der Nennbeträge beträgt 50.000,00 EUR. Im Ordner ist keine spätere Liste vorhanden. Diese Angaben stammen aus der Liste, nicht aus der Spalte zur Geschäftsführung des Registerausdrucks.

4 Finanzierungsunterlagen im Ordner

Kreditbeträge und interne Zustimmungsvorbehalte enthält der Registerausdruck nicht. Im getrennten Bankfach liegt der Kontokorrentvertrag LG-8842 vom 6. Februar 2023 über 40.000,00 EUR. Beide Geschäftsführer haben ihn unterschrieben. Die Bürgschaften sind dort als gesonderte Urkunden abgelegt. Den aktuellen Stand hat Frau Mertens in ihrem Schreiben vom 21. September 2026 beschrieben. Der Beschluss vom 20. August 2026 wurde ebenfalls ins Bankfach kopiert.''',
             'Freundliche Grüße\nBettina Kruse\nProkuristin / Übertragung am 22.09.2026 geprüft')
    document('03_Gesellschaftsvertrag_Auszug.pdf', FIRMA,
             'An die Geschäftsführung\nAusfertigungsordner Gesellschaftsvertrag',
             'Gesellschaftsvertrag Auszug aus der Büroausfertigung', 'Hannover, 22. September 2026', 'Urkunde 184/2018 vom 18.05.2018 / Seiten 3 bis 7', '''Sehr geehrte Herren,

für die Besprechung habe ich die nachstehenden Abschnitte aus der im Büro vorhandenen Ausfertigung der Gründungsurkunde von Notarin Dr. Katharina Ehlers, Hannover, übertragen. Die Nummern entsprechen den Vertragsabschnitten. Nicht übertragen sind die Gründungsbestimmungen zu Kosten und Einzahlungen sowie die Schlussbestimmungen. Nachträge sind im Ordner nicht abgelegt.

3 Stammkapital und Geschäftsanteile

Das Stammkapital beträgt 50.000,00 EUR. Lars Niemeyer übernimmt Geschäftsanteil Nr. 1 im Nennbetrag von 25.000,00 EUR. Felix Brandes übernimmt Geschäftsanteil Nr. 2 im Nennbetrag von 25.000,00 EUR. Die Einlagen sind in Geld zu leisten. Die Beteiligung am Ergebnis richtet sich nach den Nennbeträgen der Geschäftsanteile.

4 Geschäftsführung und Vertretung

Ist ein Geschäftsführer bestellt, vertritt er die Gesellschaft allein. Bei mehreren Geschäftsführern vertreten zwei Geschäftsführer gemeinsam oder ein Geschäftsführer gemeinsam mit einem Prokuristen. Die Gesellschafterversammlung kann einzelnen Geschäftsführern Einzelvertretungsbefugnis erteilen. Eine Befreiung von den Beschränkungen des Paragrafen 181 BGB bedarf eines ausdrücklichen Gesellschafterbeschlusses.

Die Geschäftsführer führen die Geschäfte nach Gesetz, Gesellschaftsvertrag und den Beschlüssen der Gesellschafter. Im Innenverhältnis bedürfen Aufnahme, Erhöhung oder Verlängerung einer Kreditlinie von mehr als 50.000,00 EUR sowie die Bestellung neuer Sicherheiten außerhalb des laufenden Liefergeschäfts der vorherigen Zustimmung der Gesellschafterversammlung. Maßgeblich ist der Gesamtbetrag der Linie nach Änderung, nicht nur der Erhöhungsbetrag. Dieser Zustimmungsvorbehalt beschränkt nicht die Vertretungsbefugnis gegenüber Dritten.

Prokura wird von der Geschäftsführung nur nach vorheriger Zustimmung der Gesellschafterversammlung erteilt. Interne Zeichnungs- oder Betragsgrenzen werden in einer gesonderten Geschäftsordnung dokumentiert. Die Erteilung von Prokura ersetzt keine Beschlussfassung über zustimmungspflichtige Geschäfte. Persönliche Verpflichtungen eines Gesellschafters werden durch einen Gesellschaftsbeschluss nicht ohne dessen eigene Erklärung begründet.

---

5 Gesellschafterversammlung und Beschlüsse

Jeder Geschäftsführer kann eine Gesellschafterversammlung einberufen. Die Einladung ist an jeden Gesellschafter unter seiner zuletzt mitgeteilten Anschrift durch eingeschriebenen Brief zu versenden. Zwischen dem Tag der Absendung und dem Tag der Versammlung müssen mindestens vierzehn volle Kalendertage liegen. Ort, Beginn und Tagesordnung sind in der Einladung anzugeben. Ergänzungen der Tagesordnung sind in gleicher Form mit einer Frist von sieben vollen Kalendertagen anzukündigen.

Die Versammlung ist beschlussfähig, wenn mindestens 75 Prozent des Stammkapitals vertreten sind. Wird diese Beteiligung nicht erreicht, ist eine neue Versammlung mit gleicher Tagesordnung einzuberufen. Diese ist ohne Rücksicht auf das vertretene Kapital beschlussfähig, wenn die zweite Einladung ausdrücklich darauf hinweist. Auch für die zweite Einladung gilt die vierzehntägige Frist.

Beschlüsse werden mit einfacher Mehrheit der abgegebenen Stimmen gefasst, soweit Gesetz oder Vertrag nichts anderes bestimmen. Je 1,00 EUR eines Geschäftsanteils gewährt eine Stimme. Enthaltungen zählen nicht als abgegebene Stimmen; Stimmengleichheit bedeutet Ablehnung. Gesetzliche Stimmverbote bleiben unberührt. Über die Versammlung ist eine Niederschrift mit Teilnehmern, Anträgen, Abstimmungsergebnissen und Beschlüssen anzufertigen und von den anwesenden Gesellschaftern zu unterzeichnen. Beschlüsse über Abschluss, Änderung und Beendigung von Geschäftsführerdienstverträgen sind der Gesellschafterversammlung vorbehalten.

6 Verfügungen und Ausscheiden

Die Abtretung und Verpfändung von Geschäftsanteilen bedürfen der Zustimmung der Gesellschafterversammlung. Dem anderen Gesellschafter steht bei einem beabsichtigten Verkauf ein Vorerwerbsrecht zu denselben Bedingungen zu. Der veräußerungswillige Gesellschafter teilt Käufer, Preis und Nebenbedingungen schriftlich mit. Die Ausübungsfrist beträgt sechs Wochen ab vollständiger Mitteilung. Die erforderliche notarielle Form bleibt unberührt.

Jeder Gesellschafter kann seine Beteiligung mit einer Frist von sechs Monaten zum Ende eines Geschäftsjahres durch eingeschriebenen Brief kündigen. Die Kündigung führt nicht zur Auflösung der Gesellschaft. Über eine Übernahme des Anteils oder eine zulässige Einziehung und die hierfür erforderlichen Erklärungen ist gesondert zu befinden. Ein bloßer Austrittswunsch bewirkt keine Auszahlung aus dem Gesellschaftsvermögen.

Kommt eine Einziehung zustande, ist die Abfindung nach dem Verkehrswert des Anteils zum maßgeblichen Ausscheidensstichtag zu bemessen. Bei Uneinigkeit wird der Wert durch einen gemeinsam zu bestellenden Wirtschaftsprüfer ermittelt. Die Zahlung soll in drei gleichen Jahresraten erfolgen, erstmals sechs Monate nach Feststellung des Werts; zwingende Kapitalerhaltungsvorschriften bleiben unberührt. Eine Einziehung gegen den Willen des Betroffenen ist nur bei Pfändung des Anteils nach erfolglosem Ablauf einer dreimonatigen Freigabefrist oder bei einem in seiner Person liegenden wichtigen Grund vorgesehen.

7 Geschäftsjahr

Geschäftsjahr ist das Kalenderjahr. Über Feststellung des Jahresabschlusses und Verwendung des Ergebnisses entscheidet die Gesellschafterversammlung. Vorschüsse auf künftige Gewinnausschüttungen bedürfen eines ausdrücklichen Beschlusses.''',
             'Mit freundlichen Grüßen\nBettina Kruse\nProkuristin / Abschrift mit der Büroausfertigung abgeglichen')
    document('04_Beschluss_2026-08-20_unterzeichnet.pdf', FIRMA,
             'An die Geschäftsführung und das Bankfach\nNiederschrift der Gesellschafterversammlung',
             'Beschluss über den Finanzierungsrahmen', 'Hannover, 20. August 2026', 'GV 2026-08 / Besprechungsraum Hägenstraße', '''Sehr geehrte Frau Kruse,

bitte nehmen Sie diese unterschriebene Niederschrift zum Gesellschaftsordner. Lars Niemeyer und Felix Brandes sind heute von 17:00 bis 17:45 Uhr persönlich anwesend. Jeder hält einen Anteil von 25.000,00 EUR. Das gesamte Stammkapital von 50.000,00 EUR ist vertreten. Beide verzichten für diese Zusammenkunft ausdrücklich auf Einladungsform und Einladungsfrist und stimmen der nachstehenden Tagesordnung zu. Herr Brandes führt die Niederschrift.

1 Betriebsmittellinie

Die Geschäftsführung wird ermächtigt, mit der Leine Gewerbebank eG eine befristete Erhöhung der bestehenden Kontokorrentlinie von 40.000,00 EUR auf insgesamt höchstens 100.000,00 EUR bis längstens 31. März 2027 zu vereinbaren. Die Erhöhung dient Materialeinkäufen und laufenden Betriebsausgaben. Der Sollzinssatz darf 9,5 Prozent jährlich und eine einmalige Bearbeitungsgebühr 750,00 EUR nicht überschreiten. Andere Finanzierungspartner werden von dieser Zustimmung nicht erfasst.

Es dürfen keine neuen Grundpfandrechte oder Sicherungsübereignungen am Fahrzeugbestand bestellt werden. Persönliche Bürgschaften bleiben einer gesonderten Entscheidung und Unterzeichnung des jeweils betroffenen Gesellschafters vorbehalten. Die Zustimmung gilt für den Abschluss durch jeden der beiden einzelvertretungsberechtigten Geschäftsführer. Sie erweitert keine Vollmacht von Frau Kruse und enthält keine Kreditzusage der Bank.

Lars Niemeyer stimmt mit 25.000 Stimmen zu. Felix Brandes stimmt mit 25.000 Stimmen zu. Gegenstimmen und Enthaltungen gibt es nicht. Der vorstehende Antrag ist einstimmig angenommen.

2 Vergütung und weitere Unterlagen

Herr Niemeyer regt für den Fall verzögerter Kundenzahlungen eine vorübergehende Absenkung beider Geschäftsführergehälter an. Herr Brandes will zunächst die August-Auswertung sehen. Es wird kein Änderungsantrag zur Abstimmung gestellt. Die bisherige feste Monatsvergütung von jeweils 6.000,00 EUR brutto bleibt in der Abrechnung bestehen. Eine Auszahlung aus einem möglichen Anteilsverkauf ist nicht Gegenstand dieser Versammlung.

Frau Kruse soll der Bank den Zwischenstand zum 31. August nachreichen. Beide Geschäftsführer erhalten eine Kopie dieser Niederschrift. Eine Änderung der Vertretungsregelungen oder der Prokura wird nicht beschlossen.''',
             'Hannover, 20.08.2026\ngez. Lars Niemeyer                    gez. Felix Brandes\nLars Niemeyer                         Felix Brandes\nBeide eigenhändigen Unterschriften sind auf der abgelegten Urschrift vorhanden.')
    document('05_Dienstvertrag_Brandes_Auszug.pdf', FIRMA,
             'Herrn Felix Brandes\nPersönlich / Geschäftsführerdienstvertrag',
             'Auszug aus dem Geschäftsführerdienstvertrag', 'Hannover, 22. September 2026', 'Vertrag vom 07.06.2018 mit Änderungsvereinbarung vom 15.12.2023', '''Sehr geehrter Herr Brandes,

auf Ihre Bitte übersende ich die nachfolgende Übertragung der derzeit im Personalsonderordner abgelegten Vergütungs- und Laufzeitregelungen. Die Änderungsvereinbarung vom 15. Dezember 2023 ersetzt ausschließlich die feste Vergütung ab Januar 2024. Eine weitere unterzeichnete Gehaltsvereinbarung befindet sich dort nicht.

1 Parteien und Aufgaben

Vertragspartner sind die Niemeyer & Brandes Gebäudetechnik GmbH und Felix Brandes. Herr Brandes übernimmt die technische Geschäftsführung, insbesondere Bauleitung, Beschaffung und fachliche Anleitung des Montagepersonals. Die organschaftliche Bestellung und die vertraglichen Beziehungen werden gesondert behandelt. Herr Brandes ist nicht als Arbeitnehmer in die Lohnliste der 15 Beschäftigten aufgenommen; seine Vergütung wird auf einem gesonderten Abrechnungskreis geführt.

2 Vergütung in der geltenden Fassung

Die Gesellschaft zahlt monatlich 6.000,00 EUR brutto, fällig am letzten Bankarbeitstag des jeweiligen Monats. Eine Tantieme oder automatische Erhöhung ist nicht vereinbart. Notwendige Geschäftsreisen werden gegen Beleg erstattet. Der zur dienstlichen und privaten Nutzung überlassene Transporter wird nach der gesonderten Fahrzeugvereinbarung abgerechnet. Ein Anspruch auf Abgeltung eines Gesellschaftsanteils ist nicht Bestandteil dieses Vertrags.

Änderungen der Vergütung bedürfen einer schriftlichen Änderungsvereinbarung auf Grundlage eines Gesellschafterbeschlusses. Das bloße Ausbleiben einer Überweisung ersetzt eine solche Vereinbarung nicht. Verzicht, Stundung oder Verrechnung sind in den abgelegten Fassungen nicht erklärt.

3 Laufzeit und Beendigung

Der Dienstvertrag ist auf unbestimmte Zeit geschlossen. Er kann mit sechs Monaten Frist zum Ende eines Kalenderjahres ordentlich gekündigt werden. Die außerordentliche Kündigung bleibt unberührt. Eine Abberufung als Geschäftsführer beendet nicht ohne Weiteres diesen Vertrag. Bei Ende der Tätigkeit sind Projektunterlagen, Schlüssel und betriebliche Zugangsmittel zurückzugeben; eine geordnete Übergabe ist mit der Gesellschaft abzustimmen.

4 Zeichnung der Fassungen

Die Fassung vom 7. Juni 2018 trägt die Unterschriften von Felix Brandes und von Lars Niemeyer für die Gesellschaft aufgrund der gesonderten Abschlussvollmacht des Gründungsbeschlusses. Die Gehaltsänderung vom 15. Dezember 2023 ist von Felix Brandes und von Bettina Kruse als hierfür durch den beigefügten Beschluss ausdrücklich bevollmächtigter Vertreterin der Gesellschaft unterschrieben. Ihre Unterschrift wird dort nicht allein auf Prokura gestützt.

Die aktuelle Vergütung von Herrn Niemeyer beträgt laut seinem getrennten Sonderordner ebenfalls 6.000,00 EUR brutto. Seine vollständige Vertragsakte habe ich diesem Auszug nicht beigefügt.''',
             'Freundliche Grüße\nBettina Kruse\nProkuristin / Übertragung aus dem Personalsonderordner')
    mail('06_Brandes_Ausstiegswunsch.eml', 'Felix Brandes <f.brandes@nb-gebaeudetechnik.de>',
         'Lars Niemeyer <l.niemeyer@nb-gebaeudetechnik.de>', '2026-09-17T20:14:00',
         'Wie es für mich weitergeht', '''Hallo Lars,

nach unserem Streit am Dienstag möchte ich meinen halben Anteil abgeben. Ich stelle mir 210.000 EUR vor und würde die operative Tätigkeit gern zum 30. November beenden. Das ist mein Verhandlungswunsch, noch keine Kündigung und kein unterschriebener Kaufvertrag. Ob Du privat kaufst oder wir jemand anderen suchen, müssen wir besprechen. Ich habe mit keinem Interessenten gesprochen.

Mir ist vor allem wichtig, dass die Bank mich aus den 40.000 EUR Bürgschaft entlässt. Eine zusätzliche Bürgschaft unterschreibe ich derzeit nicht. Ich sehe auch nicht ein, mein Gehalt zu kürzen, solange ich das Projekt Hofgarten einschließlich der Nacharbeiten zu Ende bringen soll. Wegen der zusätzlichen Abendtermine dachte ich eher an 7.000 EUR ab Oktober. Mir ist klar, dass das bislang nicht vereinbart ist.

Bitte verbuche keinen Kaufpreis als Verbindlichkeit der Firma. In Deinem Zettel gestern stand „Firma zahlt Felix aus“. So war das nicht abgesprochen. Die Maschinen sollen im Betrieb bleiben, und die Leute sollen ihre laufenden Verträge behalten. Ich verkaufe keine Werkstattinventarliste.

Für die Kundenbesprechung am Freitag bin ich verfügbar. An einer Lösung für die GmbH arbeite ich mit, auch wenn wir uns über meinen Preis nicht einig sind.

Gruß
Felix
Felix Brandes | Geschäftsführer
Niemeyer & Brandes Gebäudetechnik GmbH
Hägenstraße 38, 30559 Hannover''')
    mail('07_Niemeyer_Antwort_Gehalt.eml', 'Lars Niemeyer <l.niemeyer@nb-gebaeudetechnik.de>',
         'Felix Brandes <f.brandes@nb-gebaeudetechnik.de>', '2026-09-18T07:46:00',
         'Re: Wie es für mich weitergeht', '''Hallo Felix,

210.000 EUR kann ich privat nicht aufbringen. Ich könnte mir 120.000 EUR vorstellen, teilweise in Raten, aber meine Hausbank hat dazu noch keine Unterlagen. Ich erwarte, dass unsere gemeinsame Hausbank Deiner Entlassung nicht allein wegen eines Anteilsverkaufs zustimmt. Eine Zusage von Frau Mertens habe ich nicht.

Mit „auszahlen“ meinte ich in meinem handschriftlichen Zettel keine Verpflichtung der GmbH. Ich möchte weder Geld aus der Kasse nehmen noch den Lieferanten auf einen unklaren Verkaufstermin vertrösten. Wir sollten einen Wert anhand der tatsächlichen Zahlen besprechen. Der Augustabschluss ist noch nicht fertig, und über die 20.000 EUR vom Kunden streiten wir gerade.

Mein Vorschlag bleibt: für Oktober bis Dezember jeweils 4.000 EUR statt 6.000 EUR brutto, also für uns beide. Du hast dem nicht zugestimmt. Ich werde Frau Kruse daher keine geänderten Beträge für September geben. Deine Erhöhung auf 7.000 EUR möchte ich ebenfalls nicht freigeben. Vielleicht trennen wir die Gespräche über die künftige Arbeitsteilung von denen über den Anteil.

Wir sollten die Kanzlei für die Gesellschaft einschalten. Meine private Kaufpreisfinanzierung will ich dort nicht als Auftrag der Firma mitlaufen lassen. Am 9. Oktober kann ich ab 9 Uhr zur Gesellschafterversammlung; bitte lass uns vorher eine richtige Einladung prüfen.

Viele Grüße
Lars
Lars Niemeyer | Geschäftsführer
Niemeyer & Brandes Gebäudetechnik GmbH
Hägenstraße 38, 30559 Hannover''', reply='<06_Brandes_Ausstiegswunsch@nb-gebaeudetechnik.de>')
    document('08_Einladung_Entwurf_2026-09-23.docx', FIRMA,
             'Herrn Felix Brandes\nRutenbergstraße 21, 30559 Hannover\nAbschrift an Lars Niemeyer, Gesellschaftsordner',
             'Einladung zur Gesellschafterversammlung', 'Hannover, 23. September 2026', 'GV 2026-10 / Entwurf, noch nicht unterschrieben oder versandt', '''Sehr geehrter Herr Brandes,

hiermit lade ich Sie zur Gesellschafterversammlung am Freitag, 9. Oktober 2026, um 9:00 Uhr in den Besprechungsraum der Gesellschaft, Hägenstraße 38, 30559 Hannover, ein. Vorgesehen ist der Versand durch eingeschriebenen Brief. Ein Einlieferungsbeleg liegt noch nicht vor. Frau Kruse hat auf meine Bitte folgende Tagesordnung vorbereitet.

1 Zahlungsstand und Bankfinanzierung

Wir besprechen die Entwicklung der Kunden- und Lieferantenposten seit dem 23. September, die Septemberabrechnung sowie das Schreiben der Leine Gewerbebank vom 21. September. Zur Beratung steht, ob die Geschäftsführung den im August genehmigten Finanzierungsrahmen weiter verfolgen und welche Unterlagen sie der Bank übermitteln soll. Persönliche Bürgschaftserklärungen sind nicht Bestandteil eines Beschlusses der Gesellschaft.

2 Weiteres Vorgehen im Projekt Hofgarten

Die Geschäftsführung berichtet über den Ortstermin vom 25. September und über die geplante Inbetriebnahme. Besprochen werden sollen das weitere Zahlungsbegehren der Gesellschaft, die Zuständigkeit für die technische Korrespondenz und ein möglicher Auftrag an ein unabhängiges Messbüro. Ein Kostenangebot für eine solche Messung liegt heute noch nicht vor.

3 Geschäftsführervergütung ab Oktober

Herr Niemeyer schlägt vor, beide festen Monatsvergütungen für Oktober bis Dezember 2026 auf jeweils 4.000,00 EUR brutto zu ändern. Herr Brandes schlägt für seine eigene technische Geschäftsführertätigkeit ab Oktober 7.000,00 EUR brutto vor. Beide Vorschläge sollen getrennt beraten werden. Der Entwurf enthält weder einen gefassten Änderungsbeschluss noch eine unterzeichnete Vertragsänderung.

4 Aufgabenverteilung und Anteilsübertragung

Herr Brandes erläutert seinen Wunsch nach Abgabe seines Anteils und Beendigung seiner operativen Tätigkeit zum 30. November. Besprochen werden die noch fehlenden Bewertungsunterlagen, mögliche Gesprächspartner und eine Projektübergabe. Ein Anteilskauf, eine Einziehung oder eine Abberufung wird durch diese Einladung nicht erklärt. Ein konkreter Erwerber und ein abgestimmter Vertrag liegen nicht vor.

Bitte teilen Sie mit, ob Sie persönlich teilnehmen. Die Unterlagen zum aktuellen Bank- und Buchungsstand werden beiden Gesellschaftern vor der Sitzung auf demselben Stand zur Verfügung gestellt. Frau Kruse ist für die Erläuterung der Buchhaltung vorgesehen; eine Vollmacht zur Stimmabgabe wurde ihr nicht erteilt.''',
             'Mit freundlichen Grüßen\nLars Niemeyer\nGeschäftsführer\nUnterschrift nach Prüfung des Entwurfs vorgesehen')


def project_sources():
    document('09_Werkvertrag_Hofgarten_2026-08-03.pdf', FIRMA,
             'Hofgarten Gewerbehöfe GmbH\nFrau Maren Rösler\nBemeroder Straße 72, 30559 Hannover',
             'Auftrag für die Heizungsmodernisierung', 'Hannover, 3. August 2026', 'HG-26-041 / Angebot NB-26117 vom 28.07.2026', '''Sehr geehrte Frau Rösler,

wir bestätigen den nachstehend gemeinsam abgestimmten Werkauftrag. Auftraggeberin ist die Hofgarten Gewerbehöfe GmbH, vertreten durch Geschäftsführerin Maren Rösler. Auftragnehmerin ist die Niemeyer & Brandes Gebäudetechnik GmbH, vertreten durch Geschäftsführer Felix Brandes. Ausführungsort ist der Gewerbehof Bemeroder Straße 72, Haus 3, in Hannover.

1 Leistungen und Grenzen

Die Auftragnehmerin erneuert die Heizkreisverteilung in Haus 3, liefert und montiert sechs Pumpengruppen mit Regelmodulen sowie 24 Strangregulierventile, schließt die vorhandenen Heizstränge an und führt einen hydraulischen Abgleich einschließlich dokumentierter Einstellungen durch. Die Leistungen umfassen Demontage der Altverteiler, Entsorgung, Rohrdämmung im Technikraum und Einweisung des Hausmeisters. Das vorhandene Wärmeerzeugeraggregat bleibt bestehen.

Die Auftraggeberin stellt den trockenen, zugänglichen Technikraum und beauftragt die Herstellung der dauerhaften Stromversorgung einschließlich Potentialausgleich durch ihren Elektriker. Die Auftragnehmerin verbindet ihre Geräte ab den bereitgestellten Anschlusspunkten. Bauliche Abdichtung des Fußbodens und Umbauten an fremden Aggregaten sind nicht Teil dieses Auftrags. Das elektrische Übergabeblatt ist vor Inbetriebnahme von beiden Gewerken abzuzeichnen.

2 Preis und Zahlung

Der Pauschalpreis beträgt 120.000,00 EUR netto zuzüglich 22.800,00 EUR Umsatzsteuer, insgesamt 142.800,00 EUR brutto. Darin enthalten sind Material und Regelung mit 78.000,00 EUR netto, Montage mit 30.000,00 EUR netto sowie Abgleich, Inbetriebnahme und Einweisung mit 12.000,00 EUR netto. Zusätzliche Leistungen werden vor Ausführung unter Angabe von Preis und Termin schriftlich vereinbart.

Die erste Abschlagszahlung von 30 Prozent, 42.840,00 EUR brutto, ist nach Auftragserteilung und Zugang der entsprechenden Rechnung zahlbar. Weitere 40 Prozent, 57.120,00 EUR brutto, werden nach eingebauten Verteilern, Pumpengruppen und Anschluss der Heizstränge gegen Rechnung und Leistungsnachweis abgerechnet. Die restlichen 30 Prozent, 42.840,00 EUR brutto, werden nach Abnahme und Zugang der Schlussrechnung fällig. Jede Rechnung hat ein Zahlungsziel von vierzehn Kalendertagen ab nachgewiesenem Zugang. In der Schlussrechnung werden sämtliche erhaltenen Abschläge ausgewiesen.

---

3 Termine und Zugang

Montagebeginn ist der 17. August 2026. Die betriebsbereite Fertigstellung einschließlich Abgleich ist bis 30. September 2026 vorgesehen. Die gemeinsame Abnahme ist für den 2. Oktober 2026 um 10:00 Uhr vereinbart. Die Auftraggeberin gewährleistet werktags von 7:00 bis 17:00 Uhr Zugang; die Mietbereiche werden über Hausmeister Jens Trautmann koordiniert. Die dauerhafte Stromversorgung soll spätestens am 7. September bereitstehen. Behinderungen und Auswirkungen auf die Ausführung werden unverzüglich mitgeteilt und gemeinsam terminlich abgestimmt.

Die Auftragnehmerin misst die Betriebsgeräusche im Technikraum und kontrolliert die angrenzenden Büroräume während eines Probebetriebs. Die Parteien halten dabei die Betriebszustände und etwaige Beanstandungen fest. Die Aufnahme einzelner Heizkreise zu Testzwecken ist keine Abnahme und ersetzt das gemeinsame Protokoll nicht.

4 Abnahme und Beanstandungen

Bei der Abnahme werden die ausgeführten Leistungen besichtigt, die Einstellungen übergeben und offene Punkte in einer Liste festgehalten. Eine Unterschrift unter einen Lieferschein oder Zwischenleistungsnachweis bestätigt nur die dort bezeichneten Tatsachen. Für etwaige Mängel gelten die gesetzlichen Rechte; die VOB/B wird nicht einbezogen. Eine pauschale vertragliche Sicherheitseinbehaltsregelung wird nicht vereinbart.

Die Auftraggeberin benennt wahrgenommene Störungen möglichst mit Ort, Zeitpunkt und Betriebszustand. Die Auftragnehmerin erhält Gelegenheit, Ursache und erforderliche Arbeiten vor Ort zu prüfen. Bei akutem Wasseraustritt darf der Hausmeister den betroffenen Strang absperren und die Auftragnehmerin verständigen. Über weitergehende Eingriffe Dritter soll vorher gesprochen werden, soweit dies zeitlich möglich ist.

5 Durchführung und Abschluss

Ansprechpartner sind für die Auftraggeberin Baukoordinatorin Alina Kestner und für die Auftragnehmerin Felix Brandes. Sie koordinieren die Ausführung; Änderungen von Preis oder Leistungsumfang bedürfen der Erklärung eines hierfür bevollmächtigten Vertreters. Herr Trautmann erhält keine Befugnis, Nachträge zu bestellen. Unterlagen mit Schließ- oder Zugangsdaten werden nur an eingesetzte Mitarbeiter und beauftragte Fachunternehmen weitergegeben.

Beide Parteien haben den vollständigen vorstehenden Text gelesen und erhalten eine gleichlautende Fassung. Andere Allgemeine Geschäftsbedingungen werden nicht Vertragsbestandteil.''',
             'Hannover, 03.08.2026\ngez. Maren Rösler                       gez. Felix Brandes\nGeschäftsführerin der Auftraggeberin     Geschäftsführer der Auftragnehmerin\nÜbertragung der unterschriebenen Büroausfertigung')
    document('10_Leistungsnachweis_2026-09-02.pdf', FIRMA,
             'Hofgarten Gewerbehöfe GmbH\nFrau Alina Kestner / Baukoordination',
             'Zwischenleistungsnachweis Heizverteilung', 'Hannover, 2. September 2026', 'HG-26-041 / Begehung 14:20 bis 15:10 Uhr', '''Sehr geehrte Frau Kestner,

wir halten den heute gemeinsam besichtigten Montagezustand für die zweite Abschlagsrechnung fest. Die sechs Pumpengruppen sind eingebaut, die 24 Ventile gesetzt und die vorhandenen Heizstränge angeschlossen. Die Rohrverbindungen im Technikraum sind druckgeprüft. Die Dämmung ist dort bis auf zwei Revisionsstellen geschlossen. Das alte Verteilerstück liegt zur Abholung bereit.

Die endgültige Stromversorgung fehlt noch. Die Funktionsprobe erfolgte deshalb nur kurzzeitig über die Baustromverteilung. Die Drehzahlparameter sind Grundeinstellungen; der hydraulische Abgleich und der dokumentierte Dauerbetrieb folgen nach Herstellung des Elektroanschlusses. Heute waren die angrenzenden Büros leer. Unter regulärem Betrieb wurden dort keine Geräuschmessungen vorgenommen.

Herr Trautmann zeigt eine ältere dunkle Stelle am Sockel links der Tür. Die Stelle ist bei der heutigen Berührung trocken. Eine Feuchtemessung wurde nicht durchgeführt. Ob sie aus der früheren Anlage stammt, ist anhand dieses Termins nicht feststellbar. Herr Brandes hält ein Foto auf dem Baustellenhandy fest; Frau Kestner bittet um erneute Kontrolle im laufenden Betrieb.

Die in Abschnitt 2 des Auftrags genannten Bauteile für die zweite Abschlagsstufe sind vorhanden. Diese Feststellung bestätigt weder die vollständige Fertigstellung noch eine Abnahme. Insbesondere stehen Abgleich, Einweisung und Übergabe der Einstellwerte noch aus. Über die Leistung der kundenseitigen Elektroinstallation wird hier nichts bestätigt.

Frau Kestner erhält eine Kopie für die Rechnungsprüfung. Die Schlussbegehung bleibt für den 2. Oktober vorgemerkt. Herr Brandes kündigt den nächsten Montageeinsatz für den 8. September an; der Termin ist mit Herrn Trautmann abzustimmen.''',
             'Für die Feststellung des besichtigten Zustands:\ngez. Alina Kestner                       gez. Felix Brandes\nBaukoordination Hofgarten               Technische Geschäftsführung NB\nUnterschriften auf dem Baustellenexemplar vorhanden')
    document('11_Abschlagsrechnung_NB-260902.pdf', FIRMA,
             'Hofgarten Gewerbehöfe GmbH\nRechnungsprüfung / Bemeroder Straße 72\n30559 Hannover',
             'Zweite Abschlagsrechnung', 'Hannover, 2. September 2026', 'Rechnung NB-260902 / Debitor 10104 / HG-26-041', '''Sehr geehrte Damen und Herren,

für die vom 17. August bis 2. September 2026 ausgeführten Leistungen berechnen wir die zweite vertragliche Abschlagsstufe. Grundlage sind der Auftrag vom 3. August und der heute unterzeichnete Zwischenleistungsnachweis. Die Verteiler und Pumpengruppen sind montiert und die Heizstränge angeschlossen.

Abschlagsstufe 2: 40 Prozent des Auftragswerts von 120.000,00 EUR netto.\nNettobetrag: 48.000,00 EUR.\nUmsatzsteuer 19 Prozent: 9.120,00 EUR.\nZu zahlender Rechnungsbetrag: 57.120,00 EUR brutto.

Die erste Abschlagsrechnung NB-260803 über 36.000,00 EUR netto zuzüglich 6.840,00 EUR Umsatzsteuer, insgesamt 42.840,00 EUR brutto, wurde am 10. August 2026 vollständig bezahlt. Mit der jetzigen Rechnung sind insgesamt 84.000,00 EUR netto beziehungsweise 99.960,00 EUR brutto als Abschläge abgerechnet. Die Schlussrate von 36.000,00 EUR netto wird hier nicht berechnet.

Bitte überweisen Sie 57.120,00 EUR ohne Abzug bis 16. September 2026 auf unser bei Ihnen hinterlegtes Betriebskonto bei der Leine Gewerbebank eG, Kontoendung 8842. Zahlungsreferenz ist NB-260902 / HG-26-041. Die Rechnung geht Ihnen heute zusätzlich über das vereinbarte Rechnungsportal zu; der Eingang wurde dort um 16:42 Uhr quittiert.

Ein Sicherheitseinbehalt ist in dieser Rechnung nicht abgezogen. Für technische Rückfragen steht Herr Brandes zur Verfügung, für Buchungsfragen Frau Kruse. Unsere Umsatzsteuer-Identifikationsnummer lautet DE321784596.''',
             'Mit freundlichen Grüßen\nBettina Kruse\nProkuristin / Rechnungswesen')
    document('12_Kundenbrief_Maengel_2026-09-15.pdf', KUNDE,
             'Niemeyer & Brandes Gebäudetechnik GmbH\nHerrn Felix Brandes\nHägenstraße 38, 30559 Hannover',
             'Beanstandungen in Haus 3 und Abschlagszahlung', 'Hannover, 15. September 2026', 'HG-26-041 / Ihre Rechnung NB-260902', '''Sehr geehrter Herr Brandes,

seit dem Probebetrieb am 10. September hören unsere Mieter im Erdgeschoss ein tiefes Brummen. Besonders betroffen ist Raum 0.14, den die Steuerberatung Winterfeld nutzt. Herr Trautmann hat die Störung am 11. September um 7:40 Uhr bei Pumpenbetrieb gehört; nach Abschalten von Gruppe 2 war sie deutlich leiser. Außerdem ist die Stelle links der Technikraumtür jetzt feucht. Am 14. September lag dort ein nasses Papiertuch. Wir können nicht erkennen, ob eine Verbindung tropft oder Wasser aus dem Bodenbereich kommt.

Weiter fehlen uns lesbare Beschriftungen an zwei Reglern sowie die endgültige Liste der Einstellwerte. Ihre Monteure haben auf die noch ausstehende Elektroabnahme verwiesen. Nach Mitteilung unseres Elektrikers ist die Versorgung seit dem 9. September angeschlossen. Das gemeinsame Übergabeblatt haben wir allerdings noch nicht erhalten.

Wir bitten um Prüfung und Beseitigung der von Ihnen zu verantwortenden Beanstandungen bis 25. September 2026. Zugang ist nach telefonischer Abstimmung mit Herrn Trautmann möglich. Wenn Sie diesen Zeitraum für einzelne Arbeiten nicht einhalten können, nennen Sie uns bitte bis 18. September eine konkrete Begründung und die vorgesehenen Schritte. Wir behalten uns vor, bei fortdauerndem Wasseraustritt sofort Hilfe zu holen; derzeit ist kein aktiver Strahl sichtbar.

Von Ihrer Rechnung über 57.120,00 EUR haben wir am 14. September 37.120,00 EUR überwiesen. Weitere 20.000,00 EUR behalten wir vorerst wegen der genannten Punkte zurück. Dies ist keine auf einer Fremdrechnung beruhende Kostenschätzung. Für den Geräuschschutz hat unser Hausmeister am Telefon Beträge zwischen 8.000 und 12.000 EUR genannt; ein Angebot hat er nicht eingeholt. Die Feuchteursache ist in dieser Zahl nicht enthalten.

Der Zwischenleistungsnachweis vom 2. September war keine Abnahme. Wir möchten am vereinbarten Fertigstellungstermin festhalten und erwarten Ihren Vorschlag zum weiteren Vorgehen. Die bloße Übersendung einer korrigierten Rechnung würde unser technisches Problem nicht lösen.''',
             'Mit freundlichen Grüßen\ngez. Maren Rösler\nMaren Rösler\nGeschäftsführerin')
    document('13_Baubesprechung_2026-09-18.docx', FIRMA,
             'Frau Alina Kestner, Hofgarten Gewerbehöfe GmbH\nKopie an Herrn Jens Trautmann und die Geschäftsführung NB',
             'Besprechung im Technikraum Haus 3', 'Hannover, 18. September 2026', 'HG-26-041 / 08:30 bis 09:35 Uhr', '''Sehr geehrte Frau Kestner,

an der heutigen Besichtigung nahmen Sie, Herr Trautmann, Monteur Daniel Böttcher und ich teil. Frau Rösler war nicht anwesend. Das nachstehende Protokoll hält Beobachtungen und Absprachen fest. Über den Einbehalt haben wir keine Einigung erzielt.

1 Geräusch und Versorgung

Bei laufender Gruppe 2 war im Raum 0.14 ein Brummen wahrnehmbar. Nach Absenken des Sollwerts wurde es leiser. Herr Böttcher hat an der Wandkonsole leichte Schwingungen gespürt. Eine geeichte Schallmessung haben wir nicht vorgenommen. Ich halte eine Übertragung über die Befestigung für möglich; Herr Trautmann vermutet einen Defekt der Pumpe. Beides ist heute nicht durch Ausbau oder Messung überprüft worden.

Die Anlage lief über den endgültigen Anschluss. Die Abdeckung der Verteilung war geschlossen. Ein Prüfbericht des Elektrogewerks wurde uns nicht gezeigt. Frau Kestner will ihn bei Elektro Wende anfordern. Das elektrische Übergabeblatt ist weiterhin nicht gemeinsam unterzeichnet.

2 Feuchte und Beschriftung

Die Stelle links der Tür war am Rand feucht. Oberhalb der Stelle zeigten die sichtbaren Rohrverbindungen während des Termins keine Tropfen. Hinter der Sockelverkleidung konnten wir nicht nachsehen. Herr Trautmann meinte, der Sockel sei schon im Winter einmal dunkel gewesen; Frau Kestner kann dies nicht bestätigen. Ein Teil der Dämmung soll beim nächsten Termin geöffnet und anschließend wieder geschlossen werden.

Die beiden Regler erhalten dauerhafte Schilder. Die Einstellliste kann erst nach dem abschließenden Abgleich übergeben werden. Frau Kestner möchte die Liste trotzdem vor dem Abnahmetermin erhalten, damit der Hausmeister sie durchsehen kann.

3 Nächster Einsatz

Wir vereinbaren den 25. September um 8:30 Uhr mit Zugang zu Technikraum und Raum 0.14. NB bringt Entkopplungselemente und ein Feuchtemessgerät mit. Ein Pumpentausch wird heute nicht beauftragt. Herr Trautmann sorgt dafür, dass die Sockelverkleidung zugänglich ist. Die Firma darf noch keine kostenpflichtigen Zusatzarbeiten für Hofgarten bestellen.

Frau Kestner wird Frau Rösler über die Beobachtungen unterrichten. Sie ist nicht bereit, eine Teilfreigabe der 20.000,00 EUR zu erklären. Über Ursache, Aufwand, Terminfolgen und den verbleibenden Zahlungsbetrag wird nach dem Einsatz weiter gesprochen.''',
             'Mit freundlichen Grüßen\nFelix Brandes\nTechnische Geschäftsführung\nKenntnisnahme der Beobachtungen und des Termins: Alina Kestner, 18.09.2026, per E-Mail 12:11 Uhr bestätigt')
    plain('14_Chat_Montageteam_2026-09-21.txt', '''Niemeyer & Brandes Gebäudetechnik GmbH
Teams-Chat „Haus 3 Montage“
Export durch Daniel Böttcher am 21.09.2026, 16:58 Uhr
Zeitzone Europe/Berlin. Ausschnitt 21.09.2026, 15:41 bis 16:12 Uhr.
Teilnehmer: Daniel Böttcher, Felix Brandes, Eva Schünemann.

15:41 Daniel Böttcher: Ich habe beim Räumen noch mal Gruppe 2 laufen lassen. Auf 80 Prozent brummt es, auf 55 fast nicht. Das ist aber kein Messwert.
15:43 Felix Brandes: Freitag nicht einfach dauerhaft runterstellen. Erst Volumenstrom prüfen, sonst wird der hintere Strang nicht warm.
15:46 Daniel Böttcher: Die Konsole sitzt direkt auf dem Mauerwerk. Die Gummiteile aus dem alten Karton sind nicht verbaut. Weiß nicht mehr, ob das so mit Dir abgestimmt war.
15:49 Felix Brandes: Ich hatte mit Dir über die neue Konsole gesprochen, nicht über Weglassen der Einlage. Bring Freitag bitte die passenden Teile mit. Kein Schuldeingeständnis an den Hausmeister aus dem Chat ableiten.
15:55 Eva Schünemann: Rößler sagt, solange die Rechnung offen ist, nur Abholung gegen Vorkasse. Die Einlagen liegen bei uns im Lager. Für eine neue Pumpe hätten wir dagegen nichts frei.
16:02 Daniel Böttcher: An der Tür ist der Sockel wieder dunkel. Ich habe nicht geöffnet; Herr Trautmann musste weg. Gestern war laut ihm die Putzfirma da.
16:08 Felix Brandes: Freitag öffnen wir gemeinsam. Bitte die Stelle bis dahin nicht mit Silikon schließen. Den alten Feuchtefleck vom 2. September sehe ich auf meinem Handyfoto, aber die Dateien sind noch nicht im Projektordner.
16:12 Eva Schünemann: Zugang 25.09., 08:30 ist eingetragen. Raum 0.14 ist bis 10 Uhr frei. Frau Kestner kommt dazu.

Der Export enthält keine Fotos oder Sprachnachrichten. Beginn und Ende des gewählten Ausschnitts wurden beim Export manuell gesetzt.''')


def payment_sources():
    document('15_Lieferantenrechnung_RH-268417.pdf', LIEFERANT,
             'Niemeyer & Brandes Gebäudetechnik GmbH\nHägenstraße 38\n30559 Hannover',
             'Rechnung für Pumpengruppen und Regelung', 'Garbsen, 24. August 2026', 'RH-268417 / Kundennummer 30861 / Lieferschein LS-88217', '''Sehr geehrte Damen und Herren,

für Ihre Bestellung vom 12. August 2026 zum Projekt HG-26-041 berechnen wir die am 24. August an Ihre Werkstatt gelieferten und von Daniel Böttcher quittierten Bauteile. Es handelt sich um die abschließende Materiallieferung aus dieser Bestellung; Versandkosten fallen nicht an.

1 Lieferpositionen

6 Pumpengruppen PG-40 einschließlich Absperrung zu je 1.900,00 EUR netto: 11.400,00 EUR.\n24 Strangregulierventile SRV-25 zu je 325,00 EUR netto: 7.800,00 EUR.\n6 Regelmodule RM-8 zu je 900,00 EUR netto: 5.400,00 EUR.

Nettosumme: 24.600,00 EUR.\nUmsatzsteuer 19 Prozent: 4.674,00 EUR.\nRechnungsbetrag brutto: 29.274,00 EUR.

2 Zahlung und Lieferung

Vereinbart sind vierzehn Tage netto ohne Skonto ab Rechnungsdatum, damit Zahlung bis 7. September 2026. Bitte verwenden Sie die in Ihrem Kreditorenstamm hinterlegte Bankverbindung, Kontoendung 7318, und geben Sie RH-268417 an. Die Umsatzsteuer-Identifikationsnummer unseres Unternehmens lautet DE318426795.

Die Ware bleibt bis zur vollständigen Bezahlung nach den bei Bestellung vereinbarten Lieferbedingungen unser Eigentum. Eine gesonderte Sicherungsvereinbarung ist mit diesem Beleg nicht verbunden. Verpackung und Transportsicherung wurden bei Übergabe als vollständig bestätigt; eine Funktionsprüfung der eingebauten Anlage war nicht Teil der Anlieferung.

Bei Beanstandungen bitten wir um Nennung der betroffenen Artikel- und Seriennummer. Ihrer Buchhaltung liegt dieser Beleg seit 24. August 2026, 10:18 Uhr, per E-Mail vor. Für die Zuordnung ist Anja Lindner erreichbar.''',
             'Mit freundlichen Grüßen\nAnja Lindner\nDebitorenbuchhaltung')
    document('16_Mahnung_Roessler_2026-09-18.pdf', LIEFERANT,
             'Niemeyer & Brandes Gebäudetechnik GmbH\nFrau Bettina Kruse\nHägenstraße 38, 30559 Hannover',
             'Offene Rechnung RH-268417', 'Garbsen, 18. September 2026', 'Konto 30861 / Zahlungserinnerung vom 10.09.2026', '''Sehr geehrte Frau Kruse,

auf unsere Rechnung RH-268417 über 29.274,00 EUR mit Zahlungsziel 7. September 2026 ist bis heute keine Zahlung eingegangen. Auch nach unserer Erinnerung vom 10. September haben wir keinen Ausgleich feststellen können. Bitte lassen Sie uns den vollständigen Betrag spätestens bis Donnerstag, 24. September 2026, 12:00 Uhr zukommen und senden Sie uns den Überweisungsbeleg.

Herr Brandes erwähnte telefonisch Geräusche an einer eingebauten Pumpe. Uns fehlen dazu Artikelnummer, Seriennummer und Mess- oder Fehlerbeschreibung. Die Mitteilung allein können wir keiner konkreten gelieferten Pumpengruppe zuordnen. Eine Rücknahme oder Gutschrift ist nicht vereinbart. Bitte senden Sie technische Beanstandungen an unseren Innendienst; die Zahlungszuordnung übernimmt Frau Lindner.

Für neue Bestellungen haben wir Ihr Konto vorläufig auf Vorkasse gesetzt. Reservierte Liefertermine können wir ohne Freigabe nicht halten. Ersatzteile für den Einsatz am 25. September sind von uns bislang nicht verbindlich zugesagt. Lagerware kann nach Zahlung gegen Abholung bereitgestellt werden.

Sie hatten eine Teilzahlung in Aussicht gestellt, aber weder Betrag noch Termin genannt. Einer Ratenzahlung oder Stundung haben wir nicht zugestimmt. Die vorstehende letzte Zahlungsaufforderung ändert das ursprüngliche Zahlungsziel nicht. Verzugszinsen und weitere Kosten sind in dem hier genannten Hauptbetrag nicht enthalten; hierzu erhalten Sie gegebenenfalls eine gesonderte Berechnung.

Bitte rufen Sie mich bis 23. September an, wenn Sie eine abweichende Zahlungsfolge konkret besprechen möchten. Eine Vereinbarung müsste unsere Geschäftsführung freigeben.''',
             'Mit freundlichen Grüßen\ngez. Anja Lindner\nAnja Lindner\nDebitorenbuchhaltung')
    document('17_Bankbrief_Kreditaufstockung_2026-09-21.pdf', BANK,
             'Niemeyer & Brandes Gebäudetechnik GmbH\nHerrn Lars Niemeyer und Herrn Felix Brandes\nHägenstraße 38, 30559 Hannover',
             'Befristete Aufstockung der Betriebsmittellinie', 'Hannover, 21. September 2026', 'LG-8842 / Ihr Gespräch vom 18.09.2026', '''Sehr geehrter Herr Niemeyer, sehr geehrter Herr Brandes,

wir haben eine befristete Erhöhung Ihrer bestehenden Kontokorrentlinie von 40.000,00 EUR auf 70.000,00 EUR bis zum 31. März 2027 intern positiv vorgeprüft. Der zusätzliche Rahmen beträgt damit 30.000,00 EUR. Diese Vorprüfung ist noch keine abrufbare Zusage. Bis zur schriftlichen Bestätigung der Freischaltung gilt ausschließlich die bisherige Linie von 40.000,00 EUR.

Für die abschließende Kreditentscheidung und Vertragsausfertigung benötigen wir die betriebswirtschaftliche Auswertung einschließlich Summen- und Saldenliste August 2026, einen aktuellen Zahlungsplan, die unterschriebene Zustimmung zur Finanzierung sowie die Erläuterung der Forderung gegen Hofgarten. Bitte teilen Sie mit, ob der avisierte Gesellschafterwechsel Auswirkungen auf Geschäftsführung oder Projektabwicklung haben soll. Der übersandte Beschluss vom 20. August liegt uns vor.

Voraussetzung unserer Vorprüfung sind neue selbstschuldnerische Höchstbetragsbürgschaften von Herrn Niemeyer und Herrn Brandes über jeweils 60.000,00 EUR. Sie sollen die bisher jeweils auf 40.000,00 EUR begrenzten Bürgschaften nach Wirksamwerden ersetzen, nicht zusätzlich zu diesen bestehen. Ohne die beiderseitigen persönlichen Erklärungen kann die Erhöhung nicht freigegeben werden. Eine Entlassung eines Bürgen ist gegenwärtig nicht zugesagt.

Für den erhöhten Rahmen ist ein variabler Sollzins von derzeit 8,6 Prozent jährlich und eine einmalige Bearbeitungsgebühr von 350,00 EUR vorgesehen. Eine zusätzliche Grundschuld oder Sicherungsübereignung verlangen wir in dieser Vorprüfung nicht. Die bestehende Linie bleibt bis auf Weiteres unverändert; zusätzliche Zahlungen über den vereinbarten Rahmen werden nicht automatisch ausgeführt.

Bei vollständigem Eingang der Unterlagen bis 24. September, 12:00 Uhr, können wir eine Entscheidung für den 25. September anstreben. Einen Auszahlungstermin können wir heute nicht verbindlich bestätigen. Auf dem Geschäftskonto sind derzeit keine von uns veranlassten Sperren vermerkt. Die tatsächlich verfügbare Summe ergibt sich jeweils aus Buchungsstand, bestehenden Vormerkungen und geltender Linie.

Bitte stimmen Sie die erforderlichen Unterschriften mit uns ab. Die im Register vermerkte Einzelvertretung der Geschäftsführer bleibt von unserem Wunsch nach beiderseitiger Rückmeldung unberührt. Persönliche Bürgschaften kann Frau Kruse nicht für Sie unterzeichnen.''',
             'Mit freundlichen Grüßen\nCarolin Mertens\nFirmenkundenbetreuung\nElektronisch versandte Briefkopie')
    bank_image()
    csvfile('19_Debitoren_OP_2026-09-23.csv',
            ['Beleg', 'Kunde', 'Rechnung', 'Fällig', 'Brutto_EUR', 'Bezahlt_EUR', 'Offen_EUR', 'Buchungsnotiz'], [
                ['NB-260902', 'Hofgarten Gewerbehöfe GmbH', '02.09.2026', '16.09.2026', '57120,00', '37120,00', '20000,00', 'Einbehalt; Brief 15.09.; Schlussrate nicht abgerechnet'],
                ['NB-260910', 'Hartung Hausverwaltung GmbH', '10.09.2026', '24.09.2026', '8450,00', '0,00', '8450,00', 'Mündliche Zahlungsankündigung für 25.09.; Kruse 22.09.'],
                ['NB-260914', 'Praxisgemeinschaft Riedgarten', '14.09.2026', '28.09.2026', '9790,00', '0,00', '9790,00', 'Büro erwartet Eingang 29.09.; nicht bestätigt'],
                ['NB-260916', 'WEG Lenzstraße 14, verwaltet durch Hesse Immobilien', '16.09.2026', '30.09.2026', '3570,00', '0,00', '3570,00', 'Zahlungslauf der Verwaltung noch nicht erfragt'],
                ['NB-260921', 'Borkel Metallbau GmbH', '21.09.2026', '05.10.2026', '6545,00', '0,00', '6545,00', 'Wartung abgerechnet; Eingangstermin offen'],
            ])
    csvfile('20_Kreditoren_OP_2026-09-23.csv',
            ['Beleg', 'Lieferant', 'Rechnung', 'Fällig', 'Offen_EUR', 'Zahlungsart', 'Buchungsnotiz'], [
                ['RH-268417', 'Rößler Haustechnik Großhandel GmbH', '24.08.2026', '07.09.2026', '29274,00', 'Überweisung', 'Nicht gestundet; Mahnfrist 24.09.12 Uhr'],
                ['EW-260909', 'Elektro Wende GmbH', '09.09.2026', '25.09.2026', '4879,00', 'Überweisung', 'Werkstattanschluss Hägenstraße; nicht Kunden-Elektroauftrag'],
                ['LF-09-884', 'Leine Fahrzeugleasing GmbH', '01.09.2026', '28.09.2026', '2153,90', 'Lastschrift', 'Transporterflotte September'],
                ['EN-260930', 'Energieversorgung Südstadt GmbH', '01.09.2026', '30.09.2026', '1128,40', 'Lastschrift', 'Werkstatt und Büro September'],
                ['MI-2026-10', 'Immobiliengemeinschaft Hägenhof GbR', '01.09.2026', '01.10.2026', '3450,00', 'Dauerauftrag', 'Gewerbemiete Oktober; vollständiger Zahlbetrag'],
                ['SK-26218', 'Küster und Lohse Steuerberatung PartG mbB', '18.09.2026', '02.10.2026', '892,50', 'Überweisung', 'Finanzbuchhaltung August; Lohnbeiträge separat'],
            ])
    mail('21_Lohnbuero_September.eml', 'Miriam Lohse <m.lohse@kuester-lohse.de>',
         'Bettina Kruse <b.kruse@nb-gebaeudetechnik.de>', '2026-09-22T10:36:00',
         'Septemberabrechnung und Übermittlung der Zahlungsdatei', '''Sehr geehrte Frau Kruse,

aus dem vorläufigen Septemberlauf ergeben sich für Ihre 15 Beschäftigten 31.800,00 EUR Nettoauszahlungen. Die beiden Geschäftsführer sind in dieser Zahl nicht enthalten. Für diese ergeben sich aus jeweils unverändert 6.000,00 EUR Festgehalt brutto zusammen 7.400,00 EUR Überweisungen. Die Sachbezüge aus den Fahrzeugen sind im Abrechnungslauf berücksichtigt. Der Netto-Zahlbetrag einschließlich Geschäftsführung beträgt somit 39.200,00 EUR.

Die Überweisungen sind nach den hinterlegten Verträgen am 30. September fällig. Für die Bankdatei benötigen wir Ihre Freigabe bis 29. September, 14:00 Uhr. Das ist die technische Vorlaufzeit, keine Vorverlegung der vertraglichen Fälligkeit. Die Sozialversicherungsbeiträge September erwarten wir nach den bisherigen Daten mit 16.950,00 EUR. Sie werden am 28. September per Lastschrift eingezogen. Nachmeldungen von Stunden bis zum 24. September können diesen Betrag noch verändern. Bitte melden Sie insbesondere den Wochenenddienst von Herrn Böttcher zurück.

Die Lohnsteueranmeldung September liegt vorläufig bei 8.260,00 EUR, Zahlung am 12. Oktober. Sie ist weder im Netto-Zahlbetrag noch in den Sozialbeiträgen enthalten. Im Kreditorenexport erscheinen diese Beträge nicht, da wir sie im Lohnkreis führen. Die am 10. September fällige August-Lohnsteuer ist bereits ausgeglichen. Eine neue Gehaltsvereinbarung oder Stundung haben wir nicht erhalten und daher auch nicht gebucht.

Die August-BWA ist noch nicht freigegeben. Es fehlen die Materialabgrenzung des Projekts Hofgarten und zwei Reisekostenabrechnungen. Den vorliegenden OP-Export können Sie verwenden, aber wir können heute weder einen vollständigen Zwischenabschluss noch eine unterschriftsreife Augustauswertung an die Bank schicken. Die USt-Voranmeldung für August wurde mit Dauerfristverlängerung vorbereitet; den Zahlbetrag für Oktober teilen wir erst nach Belegabgleich mit. Er ist in den 8.260,00 EUR nicht enthalten.

Freundliche Grüße
Miriam Lohse
Steuerberaterin
Küster und Lohse Steuerberatung PartG mbB
Königstraße 36, 30175 Hannover''')
    plain('22_Anrufnotiz_Bank_2026-09-23.txt', '''Niemeyer & Brandes Gebäudetechnik GmbH
Anrufnotiz von Bettina Kruse
23.09.2026, 10:02 bis 10:18 Uhr, eingehend auf Büronummer
Gesprächspartnerin: Carolin Mertens, Leine Gewerbebank eG
Bezug: LG-8842, Schreiben vom 21.09.2026

Frau Mertens fragte nach August-BWA und Bürgschaften. Ich habe gesagt, dass die BWA bei Frau Lohse noch in Bearbeitung ist und Herr Brandes die Erhöhung seiner Bürgschaft bislang ablehnt. Frau Mertens bestätigte, dass der Zusatzrahmen von 30.000 EUR dann nicht freigeschaltet ist. Sie wird keinen früheren Bereitstellungstermin nennen, solange die Voraussetzungen fehlen. Ein Entwurf der neuen Bürgschaften sei noch nicht versandt worden.

Für die bestehende Linie gelten weiterhin 40.000 EUR. Beim Buchungsstand von minus 14.250,60 EUR stehen ohne neue Buchungen 25.749,40 EUR zur Verfügung. Das passt zum Bildschirm von 08:10 Uhr. Keine gesonderte Avalbelastung oder Kartenvormerkung sei dort zu berücksichtigen. Ein weiteres Geschäftskonto oder eine nutzbare Anlage haben wir nach unserer Buchhaltung nicht. Die Handkasse ist ein ausgeglichener Auslagenbestand ohne Bargeldbestand am heutigen Stichtag.

Ich fragte nach dem Quartalsabschluss. Frau Mertens nannte vorläufig rund 480 EUR für Zinsen und laufende Gebühren mit Belastung am 30. September; der exakte Abschluss steht noch aus. Die 350 EUR aus dem Aufstockungsbrief wären eine zusätzliche einmalige Gebühr nur bei Abschluss der neuen Vereinbarung. Ich habe nichts beauftragt und keine Überziehung über den bestehenden Rahmen angefragt.

Rückruf ist für 24.09., 11:00 Uhr vorgemerkt. Ich soll mitteilen, ob Herr Brandes eine eigene Rücksprache mit der Bank möchte. Eine Überweisung für Löhne wurde in diesem Gespräch weder eingereicht noch freigegeben.

Notiert unmittelbar nach dem Gespräch: Bettina Kruse, 23.09.2026, 10:25 Uhr.
Keine Tonaufzeichnung; Beträge zu Kosten nach mündlicher Auskunft.''')
    document('23_Buerobesprechung_2026-09-22.docx', FIRMA,
             'Lars Niemeyer und Felix Brandes\nKopie: Bettina Kruse, Buchhaltungsordner',
             'Bürobesprechung zu Zahlungen und Montage', 'Hannover, 22. September 2026', 'Notiz BK / Besprechung 16:00 bis 16:50 Uhr', '''Sehr geehrte Herren,

nachfolgend halte ich unsere heutige Bürobesprechung fest. Wir haben weder eine Gesellschafterversammlung durchgeführt noch eine neue Zahlungsvereinbarung unterschrieben. Teilgenommen haben Lars Niemeyer, Felix Brandes und ich. Frau Schünemann kam für die Einsatzplanung kurz hinzu.

1 Kundeneingänge

Ich habe heute mit Ralf Hartung telefoniert. Er will die 8.450,00 EUR aus NB-260910 am 25. September bezahlen. Eine schriftliche Bestätigung liegt nicht vor. Bei der Praxis Riedgarten steht die Rechnung über 9.790,00 EUR zum 28. September im System; aus den bisherigen Abläufen rechne ich mit dem 29. September als Eingang, habe aber niemanden erreicht. Für WEG Lenzstraße ist noch kein Zahlungslauf bekannt. Die Rechnung an Borkel ist erst am 5. Oktober fällig.

Für Hofgarten gibt es keine Zahlungsankündigung. Felix möchte am 25. September den Geräuschpunkt angehen. Er hält weder eine neue Pumpe noch die von Herrn Trautmann genannten 8.000 bis 12.000 EUR schon für erforderlich. Lars möchte, dass wir die offenen 20.000,00 EUR im Buchungsbestand lassen. Eine Gutschrift wurde nicht erstellt. Die 42.840,00 EUR Schlussrate sind noch nicht fakturiert und gehören deshalb nicht in die OP-Summe.

2 Ausgänge und Personal

Die sechs Kreditorenposten sind nach meinem heutigen Abgleich offen. Für die Liquiditätsübersicht setze ich die überfällige Rößler-Rechnung vorläufig auf den 23. September, die übrigen Beträge auf ihre bisherigen Fälligkeiten. Das ist keine bereits angelegte Zahlungsliste. Eine Teilzahlung an Rößler könnte Lars anfragen; ohne Zustimmung des Lieferanten lasse ich dessen vollen offenen Betrag stehen.

Die Lohnzahlen stammen aus Frau Lohses Mail von heute. Beide Geschäftsführer haben nur über neue Beträge gesprochen. Für September bleibt der bisherige Abrechnungslauf maßgeblich. Die Belegschaft besteht weiter aus 15 Personen; Kündigungen, Betriebsübergang oder eine Übertragung von Fahrzeugen an Felix sind nicht geplant. Die zwei Auszubildenden begleiten weiterhin feste Gesellen. Am Freitag fallen daher keine zusätzlichen Leiharbeiterkosten an.

3 Zuständigkeiten und Stand

Felix führt den Kundentermin; Lars spricht mit Rößler und mit der Kanzlei. Ich stelle OP-Listen und Kontobild bereit. Für neue Finanzierungsverträge habe ich keinen Auftrag erhalten. Unsere interne Bankmaske verlangt bei Sammelzahlungen zwei Freigaben, von denen eine von einem Geschäftsführer kommen muss. Diese technische Einstellung ist nicht mit der Einzelprokura im Register gleichzusetzen.

Die Übersicht soll bis 2. Oktober reichen. Noch nicht berücksichtigt werden können der endgültige USt-Zahlbetrag im Oktober, ein Aufwand nach Öffnung des feuchten Sockels und nicht bestellte Fremdmessungen. Für einen endgültigen Abschluss nach August fehlen noch Belege. Bis morgen möchte ich die heute besprochenen Eingangstermine als eigene Annahmen neben die OP-Daten setzen.''',
             'Mit freundlichen Grüßen\nBettina Kruse\nProkuristin\nNotiz verteilt am 22.09.2026, 17:10 Uhr')
    mail('24_Kunde_Zahlungsstand_2026-09-23.eml', 'Maren Rösler <m.roesler@hofgarten-gewerbe.de>',
         'Felix Brandes <f.brandes@nb-gebaeudetechnik.de>', '2026-09-23T11:08:00',
         'Haus 3: Termin Freitag und noch offener Betrag', '''Sehr geehrter Herr Brandes,

Frau Kestner hat mir Ihre Notiz vom 18. September gezeigt. Der Termin am Freitag um 8:30 Uhr bleibt bestehen. Herr Trautmann ist bis 11 Uhr im Haus. Bitte legen Sie uns anschließend schriftlich dar, was geöffnet, gemessen und gegebenenfalls geändert wurde. Eine bloße Reduzierung der Pumpenleistung ohne ausreichende Wärmeversorgung möchten wir nicht.

Ich habe noch keine weitere Zahlung der 20.000 EUR angewiesen. Nach dem Freitagstermin werden wir darüber sprechen. Ein festes Überweisungsdatum kann ich Ihnen daher heute nicht nennen. Die Zahlung von 37.120 EUR vom 14. September betrifft ausschließlich Ihre zweite Abschlagsrechnung, nicht eine Schlussrechnung. Sie können sie bitte entsprechend zuordnen.

Das Elektroprotokoll haben wir heute nochmals bei Herrn Wende angefordert. Ich habe die beiden Ausführungen, es sei seit dem 9. September „fertig“ und es fehle noch ein gemeinsames Übergabeblatt, bisher nicht auflösen können. Bitte verständigen Sie sich dazu direkt mit Frau Kestner. Eine zusätzliche Vergütung für Entkopplungsteile oder einen Pumpentausch haben wir nicht vereinbart.

Freundliche Grüße
Maren Rösler
Geschäftsführerin | Hofgarten Gewerbehöfe GmbH
Bemeroder Straße 72, 30559 Hannover''', cc='Alina Kestner <a.kestner@hofgarten-gewerbe.de>')
    register('25_Liquiditaetsstatus_2026-09-23.xlsx')
    mail('26_Ratenanfrage_Roessler_2026-09-23.eml', 'Lars Niemeyer <l.niemeyer@nb-gebaeudetechnik.de>',
         'Anja Lindner <a.lindner@roessler-haustechnik.de>', '2026-09-23T12:15:00',
         'RH-268417: Vorschlag zur Zahlung', '''Sehr geehrte Frau Lindner,

wie eben telefonisch angekündigt schlagen wir vor, von der offenen Rechnung RH-268417 über 29.274,00 EUR zunächst 5.000,00 EUR am 24. September und die restlichen 24.274,00 EUR am 2. Oktober zu zahlen. Bitte lassen Sie uns wissen, ob Ihre Geschäftsführung damit einverstanden ist. Eine Überweisung ist bisher nicht ausgeführt. Ich möchte Ihren Eingang nicht mit einem schon erteilten Zahlungsauftrag gleichsetzen.

Unser Kunde hat einen Teilbetrag einbehalten; wir klären die technischen Punkte am Freitag. Wir haben bei Ihnen bislang keinen eindeutig identifizierten Materialfehler angemeldet. Die Geräuschursache ist offen. Herr Brandes wird sich mit Artikel- und Seriennummer melden, falls die Untersuchung einen Hinweis auf die gelieferte Pumpe ergibt. Diese Zahlungsanfrage enthält keine Bitte um eine Gutschrift.

Die neue Banklinie steht uns noch nicht zur Verfügung. Frau Kruse hat deshalb angewiesen, Ihre gesamte Forderung im offenen Bestand zu belassen, bis eine abweichende Vereinbarung vorliegt. Für Freitag benötigen wir derzeit nur Teile aus unserem eigenen Lager. Einen weiteren Lieferauftrag wollen wir mit dieser Nachricht nicht auslösen.

Bitte bestätigen Sie auch, ob die Vorkasseeinstellung für neue Aufträge bei Annahme des Vorschlags bestehen bleibt. Sie erreichen mich heute bis 17 Uhr unter der Büronummer.

Mit freundlichen Grüßen
Lars Niemeyer
Geschäftsführer | Niemeyer & Brandes Gebäudetechnik GmbH
Hägenstraße 38, 30559 Hannover''', cc='Bettina Kruse <b.kruse@nb-gebaeudetechnik.de>')


def bank_image():
    im = Image.new('RGB', (1480, 1000), '#f3f5f6')
    draw = ImageDraw.Draw(im)
    draw.rectangle((0, 0, 1480, 82), fill='#213e45')
    draw.text((38, 24), 'Leine Gewerbebank eG', font=screen_font(28, True), fill='white')
    draw.text((1080, 28), 'Firmenkundenportal', font=screen_font(20), fill='white')
    draw.text((40, 111), 'Konten  /  Betriebskonto EUR', font=screen_font(20), fill='#4c5c61')
    draw.text((40, 154), 'Niemeyer & Brandes Gebäudetechnik GmbH', font=screen_font(30, True), fill='#202a2d')
    draw.text((40, 203), 'Konto •••• 8842     |     Abruf 23.09.2026, 08:10 Uhr     |     Bettina Kruse', font=screen_font(20), fill='#38484d')
    labels = [('Kontostand', '-14.250,60 EUR'), ('Vereinbarte Kreditlinie', '40.000,00 EUR'), ('Verfügbar', '25.749,40 EUR')]
    for index, (label, amount) in enumerate(labels):
        x = 40 + index * 480
        draw.rectangle((x, 253, x + 448, 374), fill='white', outline='#c9d1d4', width=1)
        draw.text((x + 20, 272), label, font=screen_font(20), fill='#536267')
        draw.text((x + 20, 317), amount, font=screen_font(29, True), fill='#243d42')
    draw.text((40, 410), 'Umsätze vom 11.09.2026 bis 22.09.2026', font=screen_font(23, True), fill='#263b41')
    draw.rectangle((40, 454, 1440, 500), fill='#dfe6e8')
    for x, label in [(57, 'Buchung'), (225, 'Auftraggeber / Empfänger'), (905, 'Betrag EUR'), (1165, 'Saldo EUR')]:
        draw.text((x, 467), label, font=screen_font(19, True), fill='#21363c')
    rows = [
        ('11.09.2026', 'Vortrag zum gewählten Zeitraum', '', '-39.462,60'),
        ('14.09.2026', 'Hofgarten Gewerbehöfe · NB-260902', '+37.120,00', '-2.342,60'),
        ('16.09.2026', 'Material Nord GmbH · MN-260881', '-7.760,00', '-10.102,60'),
        ('18.09.2026', 'Dammert Hausservice · NB-260904', '+8.452,00', '-1.650,60'),
        ('21.09.2026', 'Werkzeugleasing Service · WLS-0831', '-12.600,00', '-14.250,60'),
    ]
    for index, row in enumerate(rows):
        y = 513 + index * 62
        draw.rectangle((40, y - 6, 1440, y + 49), fill='white' if index % 2 == 0 else '#edf1f2')
        for x, value in zip((57, 225, 905, 1165), row):
            draw.text((x, y + 7), value, font=screen_font(20), fill='#243237')
    draw.text((40, 854), 'Keine weiteren gebuchten Umsätze am 22.09.2026. Vormerkungen: 0,00 EUR.', font=screen_font(20), fill='#36484e')
    draw.text((40, 900), 'Aktive Kreditlinie: 40.000,00 EUR. Beantragte Änderungen sind hier nicht enthalten.', font=screen_font(20), fill='#36484e')
    draw.text((40, 951), 'Ansicht gespeichert am 23.09.2026 um 08:11 Uhr', font=screen_font(17), fill='#64767c')
    im.save(register('18_Bankansicht_2026-09-23.png'))


def readme_and_rubric():
    lines = [f'# Gesellschafterkonflikt in einem Hannoveraner Handwerksbetrieb\n',
             '## 1 Vorgang\n',
             'Die Niemeyer & Brandes Gebäudetechnik GmbH beschäftigt 15 Arbeitnehmer neben zwei Geschäftsführern, die jeweils die Hälfte der Anteile halten. Einer möchte ausscheiden. Gleichzeitig sind ein Kundenprojekt, ein Rechnungseinbehalt, offene Lieferantenposten und die Septemberabrechnung zu bearbeiten. Stand der Akte: 23. September 2026, 14:00 Uhr. Mandantin ist die Gesellschaft.\n',
             '<!-- BEGIN gesamt-pdf-section (autogen) -->\n\n## 2 Downloads\n',
             'Gesamt-PDF zum Lesen, Originalformate zur Bearbeitung und Einzel-PDFs als getrennte Unterlagen. ZIPs werden mit dem Release bereitgestellt.\n',
             f'> {WARN_DE}\n>\n> {WARN_EN}\n',
             '| Fassung | Download |', '| --- | --- |',
             f'| Gesamt-PDF | [Gesamte Akte](gesamt-pdf/{SLUG}_gesamt.pdf) |',
             f'| Akten-ZIP | [Native Originaldateien](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{SLUG}.zip) |',
             f'| Einzel-PDF-ZIP | [Jede Unterlage als PDF](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{SLUG}-einzelpdfs.zip) |\n',
             '<!-- END gesamt-pdf-section (autogen) -->\n\n## 3 Bestand\n',
             '26 individuelle Quellen. Büroabschriften benennen ihren Ursprung; Entwürfe und persönliche Vorschläge sind als solche erkennbar. Die Arbeitsmappe enthält konkrete Quellen, Formeln und gespeicherte Rechenwerte. Es gibt keine Musterlösung in den Aktenstücken.\n',
             '| Nr. | Datei |', '| --- | --- |']
    lines.extend(f'| {int(n[:2])} | `{n}` |' for n in sorted(SOURCES))
    lines.extend(['\n## 4 Redaktion\n', 'Zugeordnetes Plugin: `wirtschaftsanwalt`. Die fachlichen Prüfkriterien in `rubric.yaml` sind nicht Bestandteil der Exportakte. Technische Prüfprotokolle und Vorschauen werden außerhalb des Repositorys gespeichert. Die Ursprungspersonen und Unternehmen sind für den Fall erfunden; reale öffentliche Personenprofile wurden nicht übernommen.\n'])
    (OUT / 'README.md').write_text('\n'.join(lines), encoding='utf-8')
    checks = [
        ('gesellschaftsmandat', 'human_review', 'Mandantin GmbH von den Einzelinteressen Brandes und Niemeyer abgrenzen; keine gemeinsame persönliche Vertretung stillschweigend annehmen.'),
        ('vertretung-und-innenbindung', 'human_review', 'Einzelvertretung beider Geschäftsführer, Einzelprokura, fehlende Selbstkontrahierungsbefreiung, interne Finanzierungszustimmung und technische Bankfreigabe anhand verschiedener Quellen unterscheiden.'),
        ('anteil-statt-assetdeal', 'human_review', 'Unverbindliche Preisvorstellungen, Anteilsübertragung, Satzungskündigung, Organstellung und Dienstvertrag getrennt erfassen; keinen Verkauf von Betriebsmitteln oder bereits wirksamen Austritt unterstellen.'),
        ('einberufung', 'human_review', 'Den nicht versandten Entwurf anhand der konkreten Satzungsfristen, Einladungsform, Tagesordnung und Beschlussfähigkeit prüfen; Einzelthemen und mögliche Stimmrechtsfragen ohne pauschale Vorentscheidung behandeln.'),
        ('gehaltsbestand', 'human_review', 'Bestandsvergütung September, befristeter Reduktionsvorschlag für beide Geschäftsführer und Erhöhungswunsch Brandes trennen; keine Vertragsänderung oder Lohnstundung erfinden.'),
        ('projekt-und-beweislage', 'human_review', 'Zwischenleistungsnachweis nicht als Abnahme behandeln; Geräusch, Feuchte, Beschriftung und Elektroübergabe mit ihren unterschiedlichen Beobachtungsständen bewerten; Ursache und Fremdkosten nicht als geklärt darstellen.'),
        ('abschlaege', 'human_review', '142800 EUR Bruttoauftrag, 42840 EUR bezahlter erster Abschlag, 57120 EUR zweiter Abschlag und 37120 EUR Teilzahlung zu 20000 EUR offen abstimmen; nicht fakturierte Schlussrate von 42840 EUR nicht als aktuelle OP verdoppeln.'),
        ('lieferantenstand', 'human_review', 'Ursprüngliche Fälligkeit, letzte Mahnfrist und nicht angenommene Ratenanfrage auseinanderhalten; keine Stundung, Gutschrift oder neue Lieferung als vereinbart ausgeben.'),
        ('liquiditaet-und-quellen', 'human_review', 'Kontostand minus 14250,60 EUR und bestehende Linie 40000 EUR zu 25749,40 EUR abstimmen; 30000 EUR Zusatzrahmen nicht vor Erfüllung der offenen Voraussetzungen einrechnen. Arbeitsmappe nicht mit einer fertigen Insolvenzreifeprüfung verwechseln.'),
        ('zeit-und-personal', 'human_review', '15 Arbeitnehmer und zwei gesondert abgerechnete Geschäftsführer, SV-Termin 28.09., Bankvorlauf 29.09., Lohntermin 30.09. und Lohnsteuer 12.10. auseinanderhalten; Oktober-USt und Zusatzaufwand als fehlende Informationen erkennen.'),
        ('priorisierung', 'human_review', 'Kundenortstermin, Mahnfrist und Bankrückruf konkret terminieren; fehlende Daten gezielt erheben und externe Erklärungen nur nach Mandatsumfang und Freigabe entwerfen.'),
    ]
    rubric = f'name: {SLUG}\nplugin: wirtschaftsanwalt\ndescription: Gesellschaftsmandat mit Anteilswunsch, laufendem Werkprojekt und kurzfristigen Zahlungsfragen.\nchecks:\n'
    for ident, kind, description in checks:
        rubric += f'  - id: {ident}\n    check_type: {kind}\n    description: {description}\n'
    for ident, name in [('satzung', '03_Gesellschaftsvertrag_Auszug.pdf'), ('beschluss', '04_Beschluss_2026-08-20_unterzeichnet.pdf'), ('arbeitsmappe', '25_Liquiditaetsstatus_2026-09-23.xlsx')]:
        rubric += f'  - id: {ident}\n    check_type: file_exists\n    path: {name}\n    description: Die individuelle Originalquelle liegt vor.\n'
    (OUT / 'rubric.yaml').write_text(rubric, encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--skip-workbook', action='store_true')
    args = parser.parse_args()
    (QA / 'office').mkdir(parents=True, exist_ok=True)
    corporate_sources()
    project_sources()
    payment_sources()
    rendered = render_office_batch([source for source, _ in PDF_DOCS])
    for source, target in PDF_DOCS:
        if source not in rendered:
            raise RuntimeError(f'Office-Konvertierung fehlt: {source.name}; SOFFICE konfigurieren.')
        target.write_bytes(rendered[source])
    readme_and_rubric()
    (QA / 'manifest.json').write_text(json.dumps(sorted(SOURCES), ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    if not args.skip_workbook:
        subprocess.run([node_binary(), str(ROOT / 'scripts/build-wirtschaftsanwalt-hannover-workbook.mjs')], cwd=ROOT, check=True, env=dict(os.environ, HANNOVER_AKTEN_QA=str(QA)))
    assert len(SOURCES) == 26
    print(f'{len(SOURCES)} Quellen; {len(PDF_DOCS)} Brief-/Vertrags-PDFs; keine ZIPs oder zentralen Dateien verändert.')
    print(f'Prüfdateien außerhalb der Akte: {QA}')


if __name__ == '__main__':
    main()
