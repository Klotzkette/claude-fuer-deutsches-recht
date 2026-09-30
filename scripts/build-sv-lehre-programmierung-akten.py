#!/usr/bin/env python3
"""Reproduzierbare native Originale zweier SV-Arbeitsakten; Tabellen separat mit Artifact Tool."""
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email import policy
from email.utils import format_datetime
import csv, json, re, sys, os
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[1]
QA=Path('/tmp/sozialversicherungspflicht-20260930/akten-qa')
QA.mkdir(parents=True,exist_ok=True)
WARNING='Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\nThis test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'
font_pairs=[
    (Path(os.environ['SV_TIMES_FONT']),Path(os.environ['SV_TIMES_BOLD_FONT']))
    for _ in [0] if os.environ.get('SV_TIMES_FONT') and os.environ.get('SV_TIMES_BOLD_FONT')
]+[(Path('/System/Library/Fonts/Supplemental/Times New Roman.ttf'),Path('/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf')),
   (Path('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf'),Path('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman_Bold.ttf')),
   (Path('/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf'),Path('/usr/share/fonts/truetype/liberation2/LiberationSerif-Bold.ttf')),
   (Path('/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf'),Path('/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf'))]
font_pair=next((pair for pair in font_pairs if all(p.is_file() for p in pair)),None)
if font_pair is None:raise RuntimeError('Set SV_TIMES_FONT and SV_TIMES_BOLD_FONT to installed TrueType serif fonts.')
pdfmetrics.registerFont(TTFont('TimesCase',str(font_pair[0])))
pdfmetrics.registerFont(TTFont('TimesCaseBold',str(font_pair[1])))
pdfmetrics.registerFontFamily('TimesCase',normal='TimesCase',bold='TimesCaseBold')
inventories={}

def case(slug,title,description):
    p=ROOT/'testakten'/slug;p.mkdir(parents=True,exist_ok=True)
    (p/'README.md').write_text(f'# {title}\n\n{description}\n\nDie Akte enthält persönliche Korrespondenz, unterschiedliche Vertragsstände, Abrechnungsdaten und Unterlagen zum tatsächlichen Arbeitsalltag. Eine rechtliche Lösung ist nicht enthalten.\n\n<!-- reserved-example-contacts -->\nErfundene Kontaktadressen verwenden reservierte `.example`-Domains.\n\n{WARNING}\n\n| Fassung | Download |\n| --- | --- |\n| Gesamt-PDF | [Akte vollständig](gesamt-pdf/{slug}_gesamt.pdf) |\n| Originaldateien | [Akten-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{slug}.zip) |\n| Einzel-PDFs | [Einzel-PDF-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{slug}-einzelpdfs.zip) |\n',encoding='utf-8')
    inventories[slug]=[]
    return p

def programmer():
    p=case('sozialversicherung-programmierer-leipzig','Programmierer Leipzig','Jonas „Juno“ Kern-Knörz entwickelt Software für ein Leipziger Logistikunternehmen. Ein Vermittler, wechselnde Arbeitsweisen und zwei im September 2026 zusammentreffende Verfahren hinterlassen unterschiedliche Spuren in Vertrag, Ticketsystem und Erinnerung. Aktenstand: 30. September 2026.')
    doc(p,'01_Rahmenvertrag_Nordlicht_2024.docx','Rahmenvertrag über Entwicklungsleistungen','Nordlicht Projektvermittlung GmbH · Kleiner Kielort 21 · 20144 Hamburg\nKern Software · Jonas Kern-Knörz · Karl-Heine-Straße 73 · 04229 Leipzig\nHamburg / Leipzig, 18. Januar 2024','''1 Vertragsgegenstand

Die Nordlicht Projektvermittlung GmbH, vertreten durch ihre Geschäftsführerin Anne Voigt, beauftragt Jonas Kern-Knörz, handelnd unter Kern Software, mit Leistungen der Softwareentwicklung. Die einzelnen Projekte ergeben sich aus gesonderten Leistungsabrufen. Aus diesem Rahmenvertrag folgt weder eine Verpflichtung zur Annahme eines Abrufs noch ein bestimmtes Auftragsvolumen. Ein Abruf kommt durch Bestätigung per E-Mail zustande. Ein bestehender Abruf darf nicht allein deshalb beendet werden, weil der Auftragnehmer einen weiteren Abruf ablehnt.

2 Leistungserbringung

Der Auftragnehmer organisiert seine Arbeit eigenverantwortlich und entscheidet über Arbeitsort und Arbeitszeit. Er berücksichtigt vereinbarte Übergabetermine sowie die notwendigen Abstimmungen mit dem Endkunden. Nordlicht darf fachliche Anforderungen des Endkunden weiterleiten. Eine allgemeine Verpflichtung zur Anwesenheit beim Endkunden wird hiermit nicht vereinbart. Die Nutzung von Daten des Endkunden außerhalb freigegebener Systeme richtet sich nach dessen Sicherheitsanforderungen.

Die Leistungen können durch fachlich geeignete Dritte erbracht werden. Der Auftragnehmer benennt eingesetzte Personen vorab; notwendige Zugänge und Vertraulichkeitsverpflichtungen müssen vor Beginn eingerichtet sein. Nordlicht kann den Einsatz bei konkret benannten fachlichen oder sicherheitsbezogenen Gründen ablehnen. Gegenüber Nordlicht bleibt Kern Software für die Leistung verantwortlich und rechnet etwaige Unterauftragnehmer selbst ab. Die Beauftragung weiterer Auftraggeber ist zulässig.

3 Vergütung und Nachweis

Der jeweilige Abruf legt den Stundensatz oder einen Festpreis fest. Bei Stundenvergütung stellt der Auftragnehmer monatlich Rechnung und fügt eine Übersicht nach Datum und Aufgabe bei. Die sachliche Bestätigung des Endkunden dient der Prüfung, ob die abgerechneten Leistungen dem Abruf zuzuordnen sind. Sie ist kein Nachweis über Beginn und Ende eines Arbeitstags. Nordlicht prüft eine Rechnung innerhalb von zehn Arbeitstagen. Zahlung erfolgt binnen 30 Tagen nach Eingang der prüffähigen Rechnung. Umsatzsteuer wird zusätzlich berechnet, soweit sie gesetzlich anfällt.

Kosten für normale Arbeitsmittel, Telefon und Internet sind im Honorar enthalten. Besondere Reisen werden nur nach vorheriger Kostenvereinbarung erstattet. Bei persönlicher Verhinderung entsteht kein Vergütungsanspruch für nicht erbrachte Stunden. Ein Pauschalhonorar bleibt an die darin vereinbarte Leistung gebunden. Fehler in eigenen Arbeitsergebnissen beseitigt der Auftragnehmer in angemessenem Umfang ohne zusätzliche Berechnung.

4 Rechte, Vertraulichkeit und Haftung

An individuell erstelltem und bezahltem Projektcode werden Nordlicht die zur Weitergabe an den Endkunden erforderlichen Nutzungsrechte eingeräumt. Bereits vorhandene Bibliotheken des Auftragnehmers werden im Übergabeprotokoll benannt. Ihre anderweitige Verwendung bleibt möglich, soweit keine vertraulichen Daten oder kundenspezifischen Funktionen übernommen werden. Der Auftragnehmer hält eine Berufshaftpflichtversicherung mit einer Deckung für Vermögensschäden vor und weist sie auf begründete Anfrage nach.

5 Laufzeit und Abstimmung

Der Rahmenvertrag beginnt am 18. Januar 2024. Einzelne Abrufe können eigene Laufzeiten enthalten. Der Rahmenvertrag ist mit vier Wochen Frist zum Monatsende kündbar; laufende Abrufe sind gesondert abzuwickeln. Ansprechpartnerin für Vertragsänderungen und Rechnungen ist Anne Voigt. Abweichende fachliche Wünsche des Endkunden sollen der Ansprechpartnerin zeitnah mitgeteilt werden, wenn sie Umfang oder Vergütung berühren.

Die Beteiligten verstehen die Zusammenarbeit als selbständige Tätigkeit. Kern Software führt die eigene Buchhaltung und sorgt für die eigenen Versicherungen. Die Bezeichnung als externer Entwickler wird im Projektverzeichnis beibehalten.

18. Januar 2024\nAnne Voigt · Jonas Kern-Knörz\nElektronisch ausgetauschte Vertragsfassung; Bestätigung beider Seiten liegt im E-Mail-Ordner „Nordlicht 2024“.''','Anne Voigt')
    doc(p,'02_Leistungsabruf_Migration_2024.docx','Leistungsabruf EL-24-017','Nordlicht Projektvermittlung GmbH / Kern Software\nEndkunde: Elster Logistiksysteme GmbH · Lützner Straße 169 · 04179 Leipzig\n25. Januar 2024','''1 Anlass und Ergebnis

Der Endkunde löst das bisherige System zur Erfassung von Lagerbewegungen ab. Kern Software entwickelt ein Migrationsprogramm für Artikelstammdaten und historische Bewegungen. Zum Ergebnis gehören eine wiederholbar ausführbare Importstrecke, ein Prüfbericht über abgewiesene Datensätze sowie eine technische Übergabebeschreibung. Die produktive Datenfreigabe und die Entscheidung über den Umstellungstermin verbleiben beim Endkunden.

Das erste Paket liest anonymisierte Exporte ein und erzeugt eine Fehlerliste. Es soll bis zum 15. März 2024 vorliegen. Das zweite Paket enthält die Transformation der Artikelnummern und die Verarbeitung von Nachlieferungen; vereinbarter Übergabetermin ist der 30. April 2024. Das dritte Paket umfasst Probelauf, Korrekturen und Übergabe bis zum 30. Juni 2024. Verschiebt der Endkunde die Bereitstellung seiner Probedaten um mehr als fünf Arbeitstage, stimmen die Beteiligten die Folgefristen neu ab.

2 Zusammenarbeit

Fachlicher Ansprechpartner ist Sven Pohl, Leiter Anwendungsbetrieb. Die Abstimmung findet grundsätzlich mittwochs um 14 Uhr in einer Videokonferenz von höchstens 45 Minuten statt. Weitere Termine werden bei Bedarf vereinbart. Kern Software entscheidet selbst über Architektur und Werkzeuge innerhalb der dokumentierten Schnittstellen. Eine Übernahme allgemeiner Störungen im Lagerbetrieb gehört nicht zu diesem Abruf.

Die Entwicklung erfolgt auf dem Rechner von Kern Software. Für produktionsnahe Probedaten wird ein personenbezogener Zugang zur Entwicklungsumgebung des Endkunden eingerichtet. Es werden keine unverschlüsselten Echtdaten auf lokale Datenträger kopiert. Zugang zum Leipziger Büro wird für gemeinsame Probeläufe nach Anmeldung bereitgestellt. Ein dauerhaft zugewiesener Büroarbeitsplatz ist nicht vorgesehen.

3 Vergütung und Änderungen

Es werden tatsächlich erbrachte Stunden mit 95 Euro netto vergütet. Der kalkulierte Umfang liegt bei höchstens 80 Stunden je Kalendermonat. Ein Mindestabruf besteht nicht. Bei erkennbarer Überschreitung informiert Kern Software Nordlicht vor weiteren Arbeiten. Für einzelne Pakete ist kein Festpreis vereinbart. Reisezeiten innerhalb Leipzigs werden nicht berechnet. Die Nutzungsrechte richten sich nach dem Rahmenvertrag; die bereits vorhandene Bibliothek „JunoMap“ wird im Projekt eingesetzt und bleibt im Rechteverzeichnis als Vorbestand gekennzeichnet.

Zusätzliche Schnittstellen, Schulungen oder die laufende Betreuung nach Übergabe bedürfen eines weiteren Abrufs. Eine gelegentliche Erklärung des Programmcodes im Übergabetermin ist dagegen bereits erfasst. Der Endkunde kann nach Vorlage eines Pakets konkrete Abweichungen zu den beschriebenen Anforderungen benennen. Kleine selbst verursachte Fehler werden innerhalb der vereinbarten Leistung behoben.

4 Ansprechpartner und Bestätigung

Für den Endkunden bestätigt Sven Pohl fachlichen Umfang und Zugangsbedarf. Für Nordlicht bestätigt Anne Voigt die Vergütung und den Vertragszeitraum vom 1. Februar bis zum 30. Juni 2024. Jonas Kern-Knörz bestätigt die Annahme dieses Abrufs am 26. Januar 2024. Die Bezeichnung „Projektteam Migration“ im Kalender dient ausschließlich der Terminzuordnung.

Dateistand: 25. Januar 2024; Annahmevermerk ergänzt am 26. Januar 2024.''','Anne Voigt')
    doc(p,'03_Nachtrag_Betriebsbegleitung_2025.docx','Betriebsbegleitung Elster – Anpassung für 2025','Nordlicht Projektvermittlung GmbH / Kern Software\nLeipzig / Hamburg, 12. Dezember 2024','''1 Weiterführung

Die seit dem 1. Juli 2024 mündlich erweiterte Unterstützung des Endkunden wird ab dem 1. Januar 2025 zu den folgenden Bedingungen fortgesetzt. Der ursprüngliche Migrationsabruf ist übergeben. Die aktuelle Leistung umfasst die Weiterentwicklung der Importstrecke, Unterstützung bei Datenproblemen und die Bearbeitung priorisierter Tickets aus dem Anwendungsbetrieb. Nordlicht ist weiterhin Vertragspartner und Rechnungsempfänger. Die Geschäftsführung des Endkunden hat dem Budget zugestimmt.

2 Umfang und Abstimmung

Für die Budgetplanung werden vier Leistungstage pro Woche mit bis zu acht Stunden zugrunde gelegt. Der tatsächliche Umfang kann mit dem Endkunden abweichen; eine feste Monatsvergütung oder die Zahlung nicht abgerufener Stunden wird nicht vereinbart. Der Endkunde benennt Tickets und ihre betriebliche Dringlichkeit. Über technische Umsetzung und zeitliche Schätzung entscheidet Kern Software. Überschreitet ein Ticket den vorhandenen Budgetrahmen, informiert Kern Software Sven Pohl und Nordlicht.

Für die laufende Abstimmung wird der tägliche Teamtermin um 9 Uhr genutzt. Bei Abwesenheit genügt eine kurze Nachricht mit dem Stand offener Aufgaben. Für Arbeiten an produktiven Daten und Rufbereitschaft wird eine gesonderte Freigabe verlangt. Dieser Nachtrag begründet keine nächtliche Rufbereitschaft. Die bei Störungen im Oktober 2024 versuchsweise genutzte Telefonliste wird nicht Bestandteil des Vertrags.

3 Vergütung und Arbeitsmittel

Ab dem 1. Januar 2025 beträgt der Stundensatz 105 Euro netto. Die Rechnungsstellung erfolgt unverändert monatlich anhand der bestätigten Stundenübersicht. Ein vom Endkunden ausgegebener Laptop dient dem Zugang zu geschützten Systemen und bleibt dessen Eigentum. Eigene Entwicklungswerkzeuge dürfen für lokale Arbeiten ohne Echtdaten weiter genutzt werden. Der Endkunde stellt einen Arbeitsplatz in Raum 2.14 bereit, ohne dass eine ständige Anwesenheit vereinbart wird. Die tatsächlichen Bürotage werden im Kalender abgestimmt.

4 Vertretung und Laufzeit

Die Regelung zum Einsatz fachlich geeigneter Personen aus dem Rahmenvertrag bleibt bestehen. Für den laufenden Anwendungsbetrieb muss die benannte Person mit den Datenstrukturen vertraut sein. Benötigte Zugänge sind mindestens zehn Arbeitstage vorher zu beantragen. Eine kurzfristige Übertragung eines personenbezogenen Accounts ist ausgeschlossen. Urlaubszeiten und längere Abwesenheiten werden so früh wie möglich mitgeteilt, um Ticketübernahmen innerhalb des Teams zu ermöglichen.

Die Vereinbarung läuft auf unbestimmte Zeit und kann mit vier Wochen Frist zum Monatsende gekündigt werden. Mit dieser Anpassung wird keine Abrechnung der Monate Juli bis Dezember 2024 geändert. Alle übrigen Bestimmungen des Rahmenvertrags gelten fort.

12. Dezember 2024\nAnne Voigt · Jonas Kern-Knörz\nIm Projektordner befindet sich die unterschriebene Fassung als Scan. Diese Word-Datei wurde am 13. Dezember 2024 mit der von beiden Seiten bestätigten Textfassung abgelegt.''','Anne Voigt')
    doc(p,'04_Erinnerungen_Juno_September 2026.docx','Für das Gespräch mit Herrn Wenzel','Jonas Kern-Knörz · 22. September 2026\nEigene Notizen, am 27. September nach Durchsicht des Kalenders ergänzt','''1 Wie es begonnen hat

Ich habe vorher mehrere Datenprojekte gemacht und bin durch eine Empfehlung an Nordlicht gekommen. Die Migration war mein eigenes Paket. Im Februar 2024 war ich meistens zu Hause und habe abends gearbeitet, weil tagsüber noch das Projekt für Linde Analytik lief. Sven wollte zuerst Java; ich habe Python vorgeschlagen und damit gebaut. Er wollte Ergebnisse sehen, nicht meine Anwesenheit. Eine kaputte Zuordnungstabelle habe ich an einem Sonntag repariert und nicht berechnet. Die Bibliothek hatte ich vorher schon.

2 Was danach anders war

Ab Juli 2024 gab es ständig kleine Störungen. Sven fragte, ob ich bis zur Besetzung einer Stelle helfen könne. Ich dachte an zwei Monate. Im Herbst bekam ich einen Laptop, weil der Zugang mit meinem eigenen Gerät nicht mehr ging. Seitdem nehme ich meist am Morgenmeeting teil. Manche Tickets kann nur ich bearbeiten. Aylin Bao Wendel aus dem internen Team hat mich deshalb einmal als „unseren Importkollegen“ vorgestellt. Ich korrigiere das nicht jedes Mal.

An zwei Tagen pro Woche sitze ich normalerweise im Büro. Den Platz in Raum 2.14 teilen wir mit einem Kollegen aus Dresden. Im Januar 2025 waren es wegen einer Umstellung fast vier Tage pro Woche. Ich habe keine Stempelkarte. Die Zeiten trage ich abends ein, manchmal erst freitags. Sven hat im April neun Stunden beanstandet; nach einer Diskussion blieben sechs davon bezahlt und drei unberechnet. Die Monatsrechnung enthält die korrigierten 120 Stunden.

3 Abwesenheit und andere Aufträge

Im Juli 2025 habe ich zwei Wochen Urlaub gemacht. Im Kalender musste ich „Urlaub beantragen“ anklicken. Sven sagte im Chat, die erste Woche sei schwierig; ich habe die Reise trotzdem gebucht. Aylin übernahm die Anfragen. Es gab dafür keine Bezahlung an mich. Für eine Vertretung durch Bastian hatte ich angefragt, aber die Sicherheitsschulung wäre erst nach meiner Rückkehr möglich gewesen. Bastian hat deshalb nichts gemacht. Ich habe ihm auch nichts gezahlt.

Linde Analytik hat 2025 netto 18.400 Euro gezahlt. Ich habe die Rechnungen selbst geschrieben. Die Zahlung aus dem November steht im Konto erst im Dezember. Die beigefügte Excel-Datei unterscheidet deshalb Rechnungsmonat und Eingang. Meine Akquise für einen weiteren Kunden führte zu keinem Auftrag. Ich beschäftige niemanden. Meine Krankenversicherung läuft freiwillig, für das Alter habe ich ein Wertpapierdepot und einen kleinen privaten Vertrag. Eine Entscheidung der Rentenversicherung zu meiner eigenen Beitragspflicht kenne ich nicht.

4 Die Schreiben

Den Statusantrag habe ich am 4. September selbst online abgesendet. In dem Feld nach einem bereits begonnenen Verfahren habe ich „nein“ angekreuzt, weil Sven sagte, der Prüfer komme erst am 21. September. Von dem Schreiben vom 27. August habe ich erstmals am 18. September erfahren. Ob darin mein Auftrag schon konkret genannt war, konnte ich damals nicht sehen. Die Kopie habe ich erst jetzt erhalten. Ich möchte die falsche Angabe nicht einfach stehen lassen, weiß aber nicht, was für den Beginn zählt.

Nordlicht meint, nur der Vertrag mit ihnen sei relevant. Tatsächlich haben Anne und ich im laufenden Jahr nur dreimal über Fachliches gesprochen; Tickets und Termine kommen von Sven. Ich habe nie einen Arbeitsvertrag mit Elster unterschrieben. Auch eine Zustimmung zu einer besonderen Verschiebung von Beiträgen habe ich nicht abgegeben. Bitte prüfen Sie beide Vertragsphasen und die Schreiben, bevor jemand für mich eine allgemeine Erklärung unterschreibt.

Juno.''','Jonas Kern-Knörz')
    doc(p,'05_Stoerungsprotokoll_April2025.docx','Nachbesprechung Importstörung vom 14. April','Elster Logistiksysteme GmbH · Anwendungsbetrieb\nBesprechung am 17. April 2025, 10:30–11:15 Uhr\nTeilnehmende: Sven Pohl, Aylin Bao Wendel, Jonas Kern-Knörz','''1 Ablauf

Am 14. April um 7:42 Uhr meldete das Lager doppelte Artikelbestände nach einem Import. Aylin stoppte den Folgejob. Sven rief Jonas um 8:05 Uhr auf dessen Mobiltelefon an. Jonas war auf dem Weg zu einem anderen Kundentermin und konnte ab 9:30 Uhr zugreifen. Aylin spielte bis dahin den letzten konsistenten Bestand auf die Testumgebung. Um 11:15 Uhr lag eine erste Fehlerursache vor: Die vom Lieferanten gelieferten Zeitstempel enthielten für 36 Datensätze keine Zeitzone.

Jonas ergänzte eine Prüfung und erstellte einen Korrekturlauf. Sven gab den Lauf um 13:40 Uhr frei. Die Lagerarbeit wurde um 15:10 Uhr wieder aufgenommen. Eine Erklärung an den Lieferanten verschickte Aylin. Der ursprüngliche Importcode entsprach der dokumentierten Schnittstelle; die neue Prüfung wird gleichwohl in die künftigen Standardtests aufgenommen. Eine abschließende Zuordnung des Fehleraufwands wurde in der Besprechung nicht getroffen.

2 Erreichbarkeit

Sven möchte, dass Jonas an den vereinbarten Leistungstagen im Zeitfenster 8 bis 17 Uhr auf dringende Rückfragen reagiert. Jonas weist darauf hin, dass bisher nur der Teamtermin um 9 Uhr feststeht und seine anderen Termine nicht pauschal untergeordnet werden können. Aylin schlägt vor, kritische Datenläufe am Vortag anzukündigen. Eine eigene Rufbereitschaft für Jonas wird nicht eingerichtet. Für den morgendlichen Betrieb bleibt die interne Bereitschaft zuständig.

3 Stundenfreigabe

Jonas hat für Fehleranalyse und Nachbereitung zunächst neun zusätzliche Stunden notiert. Sven hält drei Stunden Dokumentationspflege für bereits im laufenden Ticket enthalten. Jonas bietet an, diese drei Stunden diesmal nicht abzurechnen, ohne damit die Ursache der Störung zu übernehmen. Die übrigen sechs Stunden werden in der Aprilübersicht bestätigt. Mit der Korrektur ergibt sich ein Monatsumfang von 120 Stunden. Nordlicht erhält nur die freigegebene Monatsübersicht, nicht dieses Protokoll.

4 Weiteres Vorgehen

Aylin erstellt bis zum 24. April eine Checkliste für angekündigte Lieferantenänderungen. Jonas ergänzt die Zeitzonenprüfung und den Hinweis in der Übergabedokumentation. Sven klärt mit dem Einkauf, wer Lieferanteninformationen künftig an den Anwendungsbetrieb weitergibt. Für neue Aufgaben außerhalb des Importbereichs soll Sven vor der Zuweisung prüfen, ob das bestehende Budget reicht.

Protokoll: Aylin Bao Wendel, versandt am 18. April 2025. Jonas ergänzte am 22. April den Satz zur nicht übernommenen Fehlerursache. Weitere Rückmeldungen sind im Protokollordner nicht eingegangen.''','Aylin Bao Wendel')
    pdf(p,'06_Rechnung_Januar2025.pdf','Rechnung KS-2025-002','Kern Software · Jonas Kern-Knörz\nAn Nordlicht Projektvermittlung GmbH · Hamburg\nLeipzig, 3. Februar 2025','''Leistungszeitraum: 1. bis 31. Januar 2025. Betriebsbegleitung für Elster Logistiksysteme gemäß Vereinbarung vom 12. Dezember 2024. Abgerechnet werden 128 bestätigte Stunden zu 105 Euro netto. Die Stundenübersicht wurde am 31. Januar von Sven Pohl im Projektportal bestätigt.

Nettobetrag: 13.440,00 Euro. Umsatzsteuer 19 Prozent: 2.553,60 Euro. Rechnungsbetrag: 15.993,60 Euro. Bitte überweisen Sie den Rechnungsbetrag bis zum 5. März 2025 unter Angabe der Rechnungsnummer auf das bei Ihnen hinterlegte Geschäftskonto. Besondere Reisekosten oder Auslagen sind in dieser Rechnung nicht enthalten.

Die Arbeitsanteile verteilen sich auf Anpassung Importprüfung 48 Stunden, Fehleranalyse und Tickets 55 Stunden sowie Abstimmung und Dokumentation 25 Stunden. Nicht geleistete Zeiten wurden nicht berechnet. Der Umfang stimmt mit der zum Rechnungszeitpunkt bestätigten Übersicht überein.

Umsatzsteuerliche Angaben und Bankverbindung liegen dem Rechnungsempfänger im Stammdatensatz Kern Software vor. Rückfragen bitte an juno@kernsoftware.example. Vielen Dank für die Zusammenarbeit.''')
    pdf(p,'07_Kontoblatt_Februar2025.pdf','Geschäftskonto – markierter Buchungsauszug','Kern Software · Export aus der Buchhaltung am 22. September 2026\nAuszugsbereich: Februar 2025','''Buchungstag 26. Februar 2025, Wertstellung 26. Februar 2025: Gutschrift Nordlicht Projektvermittlung GmbH, 15.993,60 Euro, Verwendungszweck „KS-2025-002 EL-24-017 Januar“. Die Buchung ist der Ausgangsrechnung vom 3. Februar 2025 zugeordnet. Eine Kürzung oder Aufrechnung wurde nicht vorgenommen.

Buchungstag 28. Februar 2025: Belastung 119,00 Euro an WerkzeugCloud, Jahreslizenz Entwicklungswerkzeug. Diese Ausgabe ist im Buchhaltungskonto Softwareaufwand erfasst. Eine Weiterberechnung an Nordlicht oder Elster ist nicht erfolgt.

Der Auszug zeigt nur die für den vorgelegten Beleg markierten Buchungen. Er bildet weder den gesamten Kontoumsatz noch einen Kontostand ab. Das Programm hat die Markierung und den Exportzeitpunkt ergänzt. Die Bankdaten selbst wurden nicht verändert. Das Original des Kontoauszugs kann bei Bedarf im Onlinebanking nochmals geladen werden.

Zuordnung durch Jonas Kern-Knörz am 22. September 2026. Die beiden Buchungen betreffen unterschiedliche Geschäftsvorfälle.''')
    pdf(p,'08_Zugang_und_Arbeitsmittel.pdf','Zugangsblatt externer Projektbeteiligter','Elster Logistiksysteme GmbH · IT-Sicherheit\nAusgegeben am 7. Oktober 2024 · Person: Jonas Kern-Knörz','''Ausgegeben werden ein Laptop EL-NB-284, Netzteil und ein personenbezogener Zugang zur Entwicklungs- und Testumgebung. Das Gerät ist Eigentum von Elster Logistiksysteme. Eine Weitergabe an andere Personen und die gemeinsame Nutzung des Accounts sind untersagt. Der Account darf nur für freigegebene Projektaufgaben verwendet werden. Die Rückgabe ist bei Beendigung des Zugangsbedarfs mit dem Servicebüro zu vereinbaren.

Der Zugriff auf produktionsnahe Daten erfolgt nur über das verwaltete Gerät. Allgemeine Programmierarbeiten ohne solche Daten dürfen weiterhin auf eigenen Systemen erfolgen. Quellcode für die Elster-Projekte wird in den freigegebenen Repositories gespeichert. Sicherheitsupdates werden zentral verteilt. Die Regelungen gelten für Beschäftigte und externe Projektbeteiligte gleichermaßen.

Als Kontakt im Anwendungsbetrieb ist Sven Pohl hinterlegt. Die Zutrittskarte erlaubt den Zugang zum Büro montags bis freitags zwischen 7 und 20 Uhr. Außerhalb dieses Zeitfensters ist eine gesonderte Freigabe erforderlich. Die Zutrittsdaten werden für Sicherheitszwecke gespeichert; eine Abrechnung von Leistungsstunden auf dieser Grundlage ist nicht vorgesehen.

Empfang des Geräts bestätigt: Jonas Kern-Knörz, 7. Oktober 2024. Standortvermerk: Raum 2.14, wechselnder Arbeitsplatz. Der bestehende externe Account erhält zusätzlich eine interne Mailadresse für Ticketsystem und Besprechungseinladungen.''')
    pdf(p,'09_Betriebspruefung_Ankuendigung.pdf','Ankündigung einer Arbeitgeberprüfung','Deutsche Rentenversicherung Mitteldeutschland · Prüfdienst\nAn Elster Logistiksysteme GmbH, Geschäftsführung\n24. August 2026 · Geschäftszeichen 26-EL-7041','''Sehr geehrte Damen und Herren,

wir kündigen die Prüfung Ihres Betriebs nach Paragraf 28p SGB IV für den Zeitraum vom 1. Januar 2022 bis zum 31. Dezember 2025 an. Der vorgesehene Termin für die erste Besprechung ist der 21. September 2026 um 9 Uhr. Bitte benennen Sie eine Person, die Auskünfte zu Entgeltabrechnung und eingesetzten Fremdkräften geben kann.

Neben den Entgeltunterlagen werden die Sachkonten für Fremdleistungen sowie die dazugehörigen Rechnungen und Vertragsunterlagen benötigt. Wir werden nach Sichtung der Übersicht mitteilen, für welche Personen weitere Unterlagen erforderlich sind. Bitte halten Sie Angaben zu bereits laufenden Status- oder Einzugsstellenverfahren bereit. Eine allgemeine Übersicht über externe IT-Leistungen kann vorab elektronisch eingereicht werden.

Sollte der Termin nicht möglich sein, teilen Sie uns innerhalb einer Woche zwei alternative Termine mit. Die Übersendung von Unterlagen ersetzt nicht in jedem Fall eine weitere Nachfrage. Mit diesem Schreiben wird keine Entscheidung über eine einzelne Person oder eine Beitragshöhe getroffen.

Mit freundlichen Grüßen\nIm Auftrag\nBeate Lamm\nPrüfdienst\nEingangsvermerk Elster: 25. August 2026, 10:14 Uhr, Verwaltung.''')
    pdf(p,'10_Anforderung_Fremdleistung_August.pdf','Ergänzende Unterlagen zu IT-Fremdleistungen','Deutsche Rentenversicherung Mitteldeutschland · Prüfdienst\nAn Elster Logistiksysteme GmbH\n27. August 2026 · Geschäftszeichen 26-EL-7041','''Sehr geehrte Frau Lorenz,

nach Durchsicht der von Ihnen am 26. August übermittelten Sachkontenübersicht bitten wir um ergänzende Unterlagen zu den von Nordlicht Projektvermittlung GmbH berechneten Entwicklungsleistungen. In den Buchungstexten wird wiederholt Kern Software beziehungsweise Jonas Kern-Knörz genannt. Bitte übermitteln Sie die vorhandenen Verträge, Leistungsbeschreibungen und Stundenfreigaben für dessen Einsatz seit Februar 2024.

Wir benötigen außerdem Angaben dazu, von wem die täglichen Aufgaben erteilt wurden, wo die Leistungen tatsächlich erbracht wurden, welche Arbeitsmittel verwendet wurden und ob eine Vertretung stattgefunden hat. Falls neben Nordlicht weitere Unternehmen oder Personen an der Beauftragung beteiligt waren, erläutern Sie bitte deren Aufgaben. Die Darstellung soll Änderungen während des Prüfzeitraums erkennen lassen.

Bitte reichen Sie die Unterlagen bis zum 11. September 2026 ein. Soweit einzelne Unterlagen nur bei Nordlicht vorliegen, nennen Sie die zuständige Kontaktperson. Wir stimmen die weitere Aufklärung anschließend mit Ihnen ab. Der für den 21. September vorgesehene Besprechungstermin bleibt bestehen.

Mit freundlichen Grüßen\nIm Auftrag\nBeate Lamm\nInterner Eingangsstempel: 28. August 2026. Weiterleitung an Sven Pohl am 31. August, Betreff „Unterlagen Juno bitte zusammentragen“.''')
    pdf(p,'11_Statusantrag_Eingang.pdf','Eingangsbestätigung Onlineantrag','Deutsche Rentenversicherung Bund · Clearingstelle\nAn Jonas Kern-Knörz\n4. September 2026, 21:36 Uhr · Referenz CS-26-90418','''Ihr Antrag auf Feststellung des Erwerbsstatus wurde elektronisch übermittelt. Als Beginn der zu beurteilenden Tätigkeit haben Sie den 1. Februar 2024 angegeben. Als Vertragspartner wurde Nordlicht Projektvermittlung GmbH, Hamburg, erfasst. Im Feld „Tätigkeit bei einem Dritten“ wurde Elster Logistiksysteme GmbH, Leipzig, benannt.

Übernommene Angabe zum Verfahrensstand: „Ein Verfahren bei einer Einzugsstelle oder einem anderen Versicherungsträger wurde bisher nicht eingeleitet.“ Als Anlagen wurden Rahmenvertrag vom 18. Januar 2024, Leistungsabruf vom 25. Januar 2024 und Nachtrag vom 12. Dezember 2024 übertragen. Der Antrag beschreibt die Tätigkeit als Entwicklung und laufende Betreuung von Importsoftware.

Diese Bestätigung dokumentiert den Eingang. Sie enthält keine Aussage über die Zulässigkeit des Verfahrens oder den Erwerbsstatus. Falls Angaben ergänzt oder berichtigt werden müssen, verwenden Sie bitte die oben genannte Referenz. Weitere Unterlagen werden nach Prüfung angefordert.

Automatisch erzeugte Bestätigung. In der exportierten Fassung sind die personenbezogenen Versicherungs- und Steuerkennzeichen ausgeblendet. Die beigefügten Vertragsdateien sind im lokalen Antragsordner gespeichert.''')
    pdf(p,'12_Ticketauszug_Mai2025.pdf','Ticketliste Importteam – Monatsansicht Mai','Elster Logistiksysteme GmbH · Export am 23. September 2026\nFilter: Bearbeiter jkern, Zeitraum 01.–31.05.2025','''Der Export enthält abgeschlossene und offene Tickets aus dem Importbereich. Die Priorität wird durch Sven Pohl oder die interne Bereitschaft gesetzt. „Zugewiesen“ bezeichnet den im System eingetragenen Bearbeiter; der Export unterscheidet nicht, ob dieser die Aufgabe zuvor besprochen oder selbst übernommen hatte.

EL-4182: Lieferantentabelle erweitern. Erstzuweisung am 5. Mai durch Sven Pohl an jkern, Priorität normal, geschätzter Aufwand 12 Stunden. Kommentar Jonas vom 6. Mai: „Brauche zuerst die neuen Feldnamen; vorher kein belastbarer Termin.“ Abschluss am 12. Mai.

EL-4201: Import bricht bei Nullwert ab. Erstzuweisung am 9. Mai durch Bereitschaft, Priorität hoch. Kommentar Aylin: „Juno ist heute bei anderem Kunden, ich prüfe zuerst.“ Übernahme durch jkern am 12. Mai, Abschluss am 13. Mai.

EL-4269: Benutzerrechte Lagerverwaltung. Zuweisung an jkern durch Sven am 21. Mai. Kommentar Jonas: „Gehört nicht zur Importstrecke. Ich habe dafür weder Budget noch Kenntnis. Bitte an Infrastruktur.“ Neu zugewiesen an infrastrukturbetrieb am 22. Mai.

EL-4284: Dokumentation Monitoring. Von Jonas am 26. Mai selbst angelegt, 8 Stunden geschätzt. Sven priorisierte am 27. Mai stattdessen die Lieferantenanbindung; Ticket am Monatsende offen. Die Aufwandsangaben dieser Liste sind Schätzwerte und stimmen nicht zwingend mit abgerechneten Stunden überein.''')
    pdf(p,'13_Uebergabe_Migration_Juni2024.pdf','Übergabeprotokoll Migrationspaket','Elster Logistiksysteme GmbH / Kern Software / Nordlicht\nLeipzig, 28. Juni 2024','''Vorgeführt wurde Version 1.4.2 der Importstrecke. Der Probelauf verarbeitete 84.216 Artikel und 612.804 Bewegungsdatensätze. 73 Datensätze wurden mit dokumentierter Fehlerursache abgewiesen. Sven Pohl bestätigt, dass die vereinbarten Prüfdateien, die Ausführungsanleitung und die Liste der offenen Datenfragen übergeben wurden. Die fachliche Bereinigung der abgewiesenen Datensätze erfolgt durch den Endkunden.

Kern Software hat den Zeitzonenparameter und die Behandlung alter Artikelnummern erläutert. Die Bibliothek JunoMap ist als vorbestehender Bestandteil im Rechteverzeichnis enthalten. Die kundenspezifischen Transformationsregeln sind gesondert abgelegt. Weitere Kopien der Bibliothek bei anderen Kunden werden von dieser Übergabe nicht berührt.

Offen bleibt die Entscheidung des Endkunden, wann der produktive Erstlauf stattfinden soll. Die technische Übergabe wird dadurch nicht zurückgestellt. Die im Protokoll vom 14. Juni benannten Fehler in der Sortierfolge wurden ohne zusätzliche Berechnung behoben. Für die bisherige Migration sind keine weiteren Entwicklungspakete vereinbart.

Sven bittet um Unterstützung während der ersten Betriebswochen. Jonas stellt eine begrenzte Verfügbarkeit in Aussicht. Umfang, Abrechnung und Dauer sollen mit Nordlicht gesondert besprochen werden. Dieses Protokoll enthält dafür noch keinen neuen Auftrag.

Teilnehmende: Sven Pohl und Jonas Kern-Knörz vor Ort; Anne Voigt per Videokonferenz. Rückmeldung Sven am 1. Juli 2024: „Protokoll stimmt; bitte so ablegen.“''')
    pdf(p,'14_Rechnung_Linde_November2025.pdf','Rechnung KS-2025-024','Kern Software · Jonas Kern-Knörz\nAn Linde Analytik UG · Halle (Saale)\n20. November 2025','''Für die Erweiterung der Messdatenkonvertierung gemäß Angebot vom 2. September 2025 berechne ich den vereinbarten Festpreis von 4.400,00 Euro netto. Das Paket umfasst den Import zusätzlicher Spalten, den automatisierten Plausibilitätsbericht und eine gemeinsame Übergabesitzung. Die Übergabe erfolgte am 18. November. Die im Probelauf gefundenen Rundungsfehler wurden am 19. November ohne Mehrpreis beseitigt.

Nettobetrag 4.400,00 Euro, Umsatzsteuer 19 Prozent 836,00 Euro, Gesamtbetrag 5.236,00 Euro. Zahlung binnen 21 Tagen ohne Abzug auf das bekannte Geschäftskonto. Die Rechnung ergänzt die zuvor im März, Juni und September abgerechneten Pakete; diese bleiben unverändert.

Das Projekt wurde unabhängig von der Tätigkeit für Nordlicht und Elster durchgeführt. Ansprechpartnerin für fachliche Rückfragen ist Dr. Nele Tann. Die in diesem Projekt verwendeten Messdaten werden nicht in fremde Kundensysteme übertragen. Für eine eventuelle Erweiterung im nächsten Frühjahr erstelle ich nach Abstimmung ein neues Angebot.

Buchungsvermerk Kern Software: Zahlungseingang am 9. Dezember 2025, 5.236,00 Euro. Die Zuordnung erfolgt zur Novemberrechnung, nicht zu einer Dezemberleistung.''')
    mails=[
    ('15_Mandatsanfrage_20260922.eml','2026-09-22T20:18:00+02:00','Jonas Kern-Knörz <juno@kernsoftware.example>','Timo Wenzel <tw@wenzel-recht.example>','Statusantrag, Nordlicht und Prüfung bei Elster','''Sehr geehrter Herr Wenzel,

Frau Wendel hat mir Ihre Adresse gegeben. Ich bin seit 2024 über Nordlicht bei Elster tätig. Zuerst ging es um eine Datenmigration, inzwischen helfe ich im Betrieb. Ich habe am 4. September einen Statusantrag gestellt. Jetzt sehe ich, dass bei Elster vorher schon Unterlagen zu mir angefordert worden waren. Ich habe im Antrag verneint, dass ein Verfahren läuft, weil ich von dem konkreten Schreiben nichts wusste.

Bitte prüfen Sie, wie ich das berichtige und welche Tätigkeit überhaupt betrachtet werden muss. Mich interessiert auch, ob ich für mich selbst Beiträge nachzahlen muss, wenn die Tätigkeit weiterhin als selbständig behandelt wird. Ich habe daneben einen weiteren Kunden, aber deutlich weniger Umsatz. Ein Statusbescheid liegt nicht vor. Bis auf die private Vorsorge und meine freiwillige Krankenversicherung zahle ich nichts regelmäßig ein.

Ich hänge die Eingangsbestätigung und die nun erhaltene Anforderung vom August an. Die übrigen Unterlagen habe ich in einen Ordner exportiert. Bitte sprechen Sie zuerst mit mir; ich möchte nicht, dass Sven aus einer überraschenden Behördenanfrage von meinem Beratungsbedarf erfährt.

Freundliche Grüße
Jonas Kern-Knörz
Im Team nennt man mich Juno.''',None,('10_Anforderung_Fremdleistung_August.pdf','11_Statusantrag_Eingang.pdf')),
    ('16_Frage_Beginn_Pruefung.eml','2026-09-02T11:09:00+02:00','Jonas Kern-Knörz <juno@kernsoftware.example>','Sven Pohl <s.pohl@elsterlogistik.example>','Prüfung schon angefangen?','''Hallo Sven,

ich fülle gerade den Antrag aus. Dort steht etwas davon, ob bei einer Einzugsstelle oder einem Versicherungsträger schon ein Verfahren begonnen wurde. Du hattest letzte Woche nur gesagt, im Herbst komme jemand wegen der Lohnkonten. Geht es schon konkret um meinen Einsatz? Ich möchte da nichts Falsches angeben.

Wenn ihr meinen Vertrag weiterreichen wollt, schick mir bitte vorher, welchen Stand ihr habt. Die erste Leistungsbeschreibung passt zum heutigen Tagesgeschäft nur noch teilweise. Anne hat mir gesagt, der spätere Nachtrag liege auch bei euch. Ich habe keine eigene Entscheidung aus einem früheren Auftrag.

Ich bin heute ab 14 Uhr bei Linde. Am besten antworte per Mail, dann habe ich das direkt beim Formular. Danke!
Juno''',None,()),
    ('17_Antwort_Sven_Pruefbeginn.eml','2026-09-02T12:26:00+02:00','Sven Pohl <s.pohl@elsterlogistik.example>','Jonas Kern-Knörz <juno@kernsoftware.example>','Re: Prüfung schon angefangen?','''Hi Juno,

der Termin ist erst am 21. September. Angefangen hat hier noch niemand, soweit ich weiß. Die Verwaltung sammelt Belege. Ich habe die Bitte bekommen, Verträge und ein paar Stundenlisten zusammenzustellen; ob die Prüferin die schon gesehen hat, weiß ich nicht. Ich würde das Feld deshalb aktuell mit nein beantworten, aber ich bin bei diesen Anträgen kein Experte.

Ich habe beide Verträge und den Nachtrag. Die Annahme, du wärst seit Februar 2024 immer an vier Tagen hier, wäre jedenfalls falsch. Das wurde erst später mehr. Für den ersten Zeitraum habe ich nur unsere Mittwochstermine im Kalender.

Lass uns Freitag kurz sprechen. Die Anforderung der Verwaltung habe ich gerade nicht zur Hand.
Sven''','16_Frage_Beginn_Pruefung.eml',()),
    ('18_Nordlicht_Vertretung.eml','2025-06-11T15:04:00+02:00','Anne Voigt <anne.voigt@nordlicht-projekte.example>','Jonas Kern-Knörz <juno@kernsoftware.example>','Bastian / Zugang für Juli','''Hallo Jonas,

fachlich hat Sven nichts gegen Bastian. Die Schulung für neue Zugänge findet bei Elster aber erst am 23. Juli statt. Vorher kann er keine produktionsnahen Daten sehen. Mit deinem Account darf er ausdrücklich nicht arbeiten. Für die zwei Wochen lohnt sich ein eingeschränkter Zugang zur Dokumentation wahrscheinlich nicht.

Bitte kläre mit Sven, welche Tickets Aylin während deiner Abwesenheit übernehmen kann. Wenn Bastian später ein eigenes Paket machen soll, brauche ich seine Vertraulichkeitserklärung und eine kurze Leistungsbeschreibung. Sein Honorar würdest du mit ihm vereinbaren; Nordlicht würde weiterhin deine Rechnung erhalten. Für Juli ist das damit erst einmal vom Tisch.

Ich habe deine angekündigten zwei Wochen ohne Leistungsstunden in der Vorschau vermerkt. Einen Urlaub bezahlen wir natürlich nicht. Bitte gib die tatsächlich geleisteten Stunden für die übrigen Wochen wie üblich an.

Viele Grüße
Anne''',None,()),
    ('19_Nachfrage_Clearingstelle.eml','2026-09-08T10:12:00+02:00','Clearingstelle <eingang@clearingstelle-post.example>','Jonas Kern-Knörz <juno@kernsoftware.example>','CS-26-90418 – ergänzende Angaben','''Sehr geehrter Herr Kern-Knörz,

zu Ihrem am 4. September eingegangenen Antrag benötigen wir ergänzende Angaben. Bitte beschreiben Sie die Tätigkeit getrennt nach dem ursprünglichen Migrationsprojekt und der späteren Betriebsbegleitung. Erläutern Sie außerdem, welche Aufgaben Nordlicht Projektvermittlung GmbH bei der Durchführung übernimmt und wer beim Endkunden Aufgaben, Prioritäten und Termine vorgibt.

Bitte teilen Sie uns mit, ob Ihnen inzwischen ein anderes Verfahren bekannt geworden ist, das dieselbe Tätigkeit betrifft. Falls Ihnen hierzu Schreiben vorliegen, reichen Sie diese mit Eingangsdatum ein. Dies gilt auch für eine Prüfung beim Endkunden. Über die weitere Bearbeitung können wir erst nach Prüfung der Angaben entscheiden.

Wir bitten um Rückmeldung bis zum 30. September 2026 unter Angabe der Referenz. Diese Nachricht enthält keine Entscheidung zum Erwerbsstatus. Unterlagen können über den bereits verwendeten sicheren Zugang eingereicht werden.

Mit freundlichen Grüßen
Im Auftrag
Carolin Voss''',None,()),
    ('20_Weiterleitung_Anforderung.eml','2026-09-18T16:38:00+02:00','Aylin Bao Wendel <a.wendel@elsterlogistik.example>','Jonas Kern-Knörz <juno@kernsoftware.example>','Das Augustschreiben zu deinen Unterlagen','''Hi Juno,

Sven hat mich gebeten, dir die Anforderung zu schicken, weil du nach dem Datum gefragt hattest. Der Anhang kam laut Verwaltung am 28. August an. Ich hatte die Mail von dort am 31. August im Projektordner abgelegt. Ich weiß nicht, ob Sven sie vor seiner Antwort an dich komplett gelesen hat. Er war die Woche wegen der Lagerumstellung dauernd unterwegs.

Die reine Terminankündigung ist noch ein anderes Schreiben. Im Anhang hier steht dein Name. Am21. September soll jetzt zunächst die Buchhaltung dran sein; für unsere Unterlagen ist vielleicht ein weiterer Termin nötig. Verlass dich bitte nicht auf meinen Kalenderstand, ich koordiniere die Prüfung nicht.

Ich habe dir außerdem den Ticketauszug zugesagt. Der kommt nächste Woche, weil ich die Kundendaten ausblenden muss. Die Kommentare zu unserer Mai-Störung bleiben drin.

Viele Grüße
Aylin''',None,('10_Anforderung_Fremdleistung_August.pdf',)),
    ('21_Nordlicht_Vertragsauffassung.eml','2026-09-21T09:57:00+02:00','Anne Voigt <anne.voigt@nordlicht-projekte.example>','Jonas Kern-Knörz <juno@kernsoftware.example>','Unterlagen und Stand des Antrags','''Hallo Jonas,

ich habe von der Prüfung bei Elster inzwischen gehört. Unser Vertragspartner bist du, Elster hat einen eigenen Vertrag mit uns. Aus meiner Sicht muss das in jeder Antwort klar stehen. Wir vereinbaren deine Vergütung, tragen gegenüber Elster das Ausfallrisiko der Rechnung und haben die Vertretung ausdrücklich ermöglicht. Dass Sven fachliche Tickets bespricht, war von Anfang an zu erwarten.

Bitte schreib aber nicht, dass wir täglich die Arbeit koordinieren. Das tun wir nicht. Ich bekomme Monatsübersichten und werde bei Budgetänderungen eingeschaltet. Wie oft du im Büro bist, weiß ich nur aus unseren gelegentlichen Gesprächen. Für den späteren Ablauf kannst du mehr sagen als ich.

Schick mir bitte vor einer abgestimmten Darstellung den Entwurf. Wir werden keine Erklärung abgeben, nach der alle früheren Einsätze identisch waren. Für 2024 waren Umfang und Aufgaben deutlich anders. Die bisherigen Rechnungen bleiben bis zu einer belastbaren Klärung normal fällig.

Viele Grüße
Anne''',None,()),
    ('22_Fristverlaengerung_Anfrage.eml','2026-09-29T17:22:00+02:00','Jonas Kern-Knörz <juno@kernsoftware.example>','Clearingstelle <eingang@clearingstelle-post.example>','CS-26-90418 – Unterlagen und Bitte um kurze Fristverlängerung','''Sehr geehrte Frau Voss,

ich bitte darum, die Frist für meine ergänzende Darstellung bis zum 9. Oktober 2026 zu verlängern. Ich habe erst nach Antragstellung erfahren, dass beim Endkunden bereits am 27. August Unterlagen zu meinem Einsatz angefordert wurden. Eine Kopie habe ich am 18. September erhalten. Ich werde das Schreiben und die genaue zeitliche Darstellung nachreichen. Meine frühere Angabe beruhte auf der Auskunft, die Prüfung beginne erst am 21. September.

Die Vertragsunterlagen aus dem Antrag bleiben vollständig. Zusätzlich stelle ich die Abläufe der Migration und der späteren Unterstützung getrennt zusammen. Mein beratender Rechtsanwalt erhält morgen den vollständigen Ordner. Ich bitte um kurze Bestätigung, ob die Fristverlängerung möglich ist.

Mit freundlichen Grüßen
Jonas Kern-Knörz

Lokaler Versandvermerk: aus dem Postausgang am 29. September, 17:22 Uhr übertragen. Bis zum Export am30. September,08:30Uhr ist keine Antwort im Posteingang enthalten.''',None,())]
    for row in mails:mail(p,*row)
    raw(p,'23_Chat_Urlaub_Juli2025.txt','''Teams-Export · Gespräch Sven Pohl / Juno Kern-Knörz · 03.–06.06.2025
03.06.2025 10:14 Juno: Ich bin vom 7. bis 18. Juli weg. Wo trage ich das ein, damit keine Tickets bei mir hängen bleiben?
03.06.2025 10:20 Sven: Im Abwesenheitskalender, Typ Urlaub. Die erste Woche ist wegen L3 eigentlich schlecht.
03.06.2025 10:25 Juno: Reise ist gebucht. Das ist die einzige gemeinsame Zeit mit meiner Partnerin. Ich kann vorher dokumentieren.
03.06.2025 10:28 Sven: Ich rede mit Aylin. Dann bitte bis Ende Juni Übergabe. Hast du eine Vertretung?
03.06.2025 10:34 Juno: Bastian könnte, wenn ihr den Zugang macht. Ich frage Anne wegen der Papiere.
06.06.2025 09:02 Sven: Aylin deckt die Tage erst mal ab. Habe den Antrag im Kalender freigegeben.
06.06.2025 09:07 Juno: Danke. Nur damit es klar ist: Für die zwei Wochen schreibe ich keine Stunden. Erreichbar bin ich auf der Reise auch nicht zuverlässig.
06.06.2025 09:10 Sven: Verstanden. Dann bitte automatische Abwesenheitsantwort setzen.
Exportiert durch Jonas am 23.09.2026. Reaktionen und Profilbilder nicht enthalten.''')
    raw(p,'24_Chat_Standup_April2025.txt','''Teams-Kanal Importteam · Auszug 07.–09.04.2025
07.04.2025 08:46 Aylin: Ich moderiere heute. Bitte alle Tickets vorher aktualisieren.
07.04.2025 08:49 Juno: Bin beim anderen Kunden und heute erst 11 Uhr online. EL4010 läuft im Test, Link steht im Ticket.
07.04.2025 08:52 Sven: Juno, bei vier eingeplanten Tagen brauchen wir dich wenigstens im Neun-Uhr-Termin. Sonst verteilen wir ins Leere.
07.04.2025 08:56 Juno: Montagvormittag war diese Woche nie angeboten. Kalender steht seit Freitag so. Mein Angebot ist Di–Fr.
07.04.2025 09:04 Aylin: Dann ziehen wir die Abstimmung zur Importprüfung auf 11:30. Für heute reicht der Ticketstand.
08.04.2025 09:31 Sven: Nachtrag gestern: Bitte Abweichungen vom üblichen Rhythmus auch hier schreiben. Kalender allein geht unter.
08.04.2025 10:03 Juno: Kann ich machen. Dafür bitte neue Aufgaben erst mit mir abstimmen und nicht einfach meinen Namen eintragen.
09.04.2025 14:06 Aylin: EL4022 habe ich dir vorläufig zugewiesen, weil nur du die alte Mappingdatei kennst. Können wir morgen 5 Minuten schauen?
09.04.2025 14:11 Juno: Ja. Die Schätzung mache ich nach Sichtung, nicht vorher.
Technischer Export: Textnachrichten einschließlich nachträglicher Änderungen, ohne gelöschte Inhalte.''')
    raw(p,'25_Telefonnotiz_Angebot2024.txt','''Jonas Kern-Knörz · Notiz vom16.01.2024 nach Gespräch mit Anne Voigt
Empfehlung über Nele. Elster braucht kein vollständiges Team, sondern jemanden für alte Daten. Ich soll einen kleinen Probelauf zeigen, bevor sie den Auftrag bestätigen. Datenschnittstelle ist dokumentiert, Datenqualität unsicher. Anne sagt: voraussichtlich fünf Monate, 80 Stunden/Monat als Deckel. Wenn sie weniger Daten liefern, entsteht kein Mindesthonorar.
Ich: 95 EUR pro Stunde, keine Vollzeitreservierung. Bei Linde läuft bis März noch ein Paket. Arbeiten remote, einmal wöchentlich sprechen okay. Python statt Java klären. Eigene Bibliothek bleibt meine. Keine allgemeine Hotline nach Abschluss.
Anne: Elster will einmal sehen, dass ich mit Problemfällen umgehen kann. Ein kostenloser halber Tag Demo wäre gut. Ich: maximal zwei Stunden mit anonymisierten Beispielen, keine echte Entwicklung. Zugestimmt.
Offen: Daten bis 24. Januar; Termin Sven am 22. Januar 14 Uhr; Versicherungsnachweis aus Ordner holen. Anreise zum Probelauf selbst tragen. Bei Folgeauftrag Preise neu besprechen.
Ergänzung am 26. Januar: Abruf angenommen. Demo zwei Stunden nicht berechnet. Python freigegeben. Migrationsstart 1. Februar.''')
    hours=[128,136,144,120,136,128,96,112,144,152,128,88]
    paydates=['2025-02-26','2025-03-26','2025-04-25','2025-05-26','2025-06-26','2025-07-25','2025-08-26','2025-09-26','2025-10-27','2025-11-26','2025-12-23','2026-01-26']
    csvfile(p,'26_Stunden_und_Rechnungen_2025.csv',[['Leistungsmonat','Freigegebene_Stunden','Satz_netto_EUR','Netto_EUR','USt_EUR','Brutto_EUR','Zahlungseingang','Hinweis']]+[[f'2025-{i:02}',h,105,h*105,round(h*105*.19,2),round(h*105*1.19,2),paydates[i-1],('April um 3 Stunden gekürzt' if i==4 else 'Abwesenheit 7.–18.07.' if i==7 else '')] for i,h in enumerate(hours,1)])
    return p

def register(p,name,title):
    inventories[p.name].append({'file':name,'title':title})

def paragraphs(text):
    return [x.strip() for x in text.strip().split('\n\n') if x.strip()]

def doc(p,name,title,meta,text,author):
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2.0);s.bottom_margin=Cm(2.0);s.left_margin=Cm(2.3);s.right_margin=Cm(2.3)
    for st in ['Normal','Title','Heading 1','Heading 2']:
        style=d.styles[st];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0);style.font.size=Pt(11)
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():del fonts.attrib[attr]
        for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
        for border in list(style.element.xpath('.//w:pBdr')):border.getparent().remove(border)
        for spacing in list(style.element.xpath('.//w:contextualSpacing')):spacing.getparent().remove(spacing)
        style.paragraph_format.space_after=Pt(7)
    d.styles['Title'].font.size=Pt(14);d.styles['Title'].font.bold=True
    d.styles['Heading 1'].font.size=Pt(11);d.styles['Heading 1'].font.bold=True
    d.styles['Heading 1'].paragraph_format.space_before=Pt(11)
    d.styles['Normal'].paragraph_format.line_spacing=1.08
    if name.startswith('04_Erinnerungen_Juno'):
        d.styles['Normal'].paragraph_format.space_after=Pt(4)
    d.add_paragraph(title,'Title');d.add_paragraph(meta)
    for item in paragraphs(text):
        if re.match(r'^\d+(?:\.\d+)*\s+[^\n]+$',item) and len(item)<110:d.add_paragraph(item,'Heading 1')
        else:d.add_paragraph(item)
    foot=s.footer.paragraphs[0];foot.alignment=2
    run=foot.add_run('Seite ');run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    d.core_properties.author=author;d.core_properties.title=title;d.core_properties.language='de-DE'
    d.save(p/name);register(p,name,title)

def pdf(p,name,title,meta,text,table=None):
    body=ParagraphStyle('Body',fontName='TimesCase',fontSize=11,leading=14.2,spaceAfter=8)
    heading=ParagraphStyle('Heading',parent=body,fontName='TimesCaseBold',fontSize=15,leading=18,spaceAfter=15)
    small=ParagraphStyle('Meta',parent=body,fontSize=9,leading=11,spaceAfter=15)
    story=[Paragraph(escape(title),heading),Paragraph(escape(meta).replace('\n','<br/>'),small)]
    for item in paragraphs(text):story.append(Paragraph(escape(item).replace('\n','<br/>'),body))
    if table:
        rows=[[Paragraph(escape(str(c)),ParagraphStyle('Cell',parent=body,fontSize=9,leading=11,spaceAfter=2)) for c in row] for row in table]
        t=Table(rows,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8E8E8')),('LINEBELOW',(0,0),(-1,0),0.5,colors.black),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.append(t)
    def page(c,d):
        c.setFont('TimesCase',9);c.drawRightString(A4[0]-55,33,f'{d.page}')
    SimpleDocTemplate(str(p/name),pagesize=A4,rightMargin=60,leftMargin=60,topMargin=48,bottomMargin=48,title=title,author=meta.split('\n')[0]).build(story,onFirstPage=page,onLaterPages=page)
    register(p,name,title)

def mail(p,name,date,frm,to,subject,text,reply=None,attachments=()):
    m=EmailMessage(policy=policy.SMTP);m['From']=frm;m['To']=to;m['Date']=format_datetime(datetime.fromisoformat(date));m['Subject']=subject;m['Message-ID']=f'<{p.name}.{name[:-4]}@aktenpost.example>'
    if reply:m['In-Reply-To']=f'<{p.name}.{reply[:-4]}@aktenpost.example>';m['References']=m['In-Reply-To']
    m.set_content(text.strip()+'\n',charset='utf-8')
    for filename in attachments:
        source=p/filename
        mt='application';sub='pdf' if source.suffix=='.pdf' else 'vnd.openxmlformats-officedocument.wordprocessingml.document'
        m.add_attachment(source.read_bytes(),maintype=mt,subtype=sub,filename=filename)
    (p/name).write_bytes(m.as_bytes());register(p,name,subject)

def raw(p,name,text):
    (p/name).write_text(text.strip()+'\n',encoding='utf-8');register(p,name,name)

def csvfile(p,name,rows):
    with (p/name).open('w',encoding='utf-8',newline='') as f:csv.writer(f,delimiter=';').writerows(rows)
    register(p,name,name)

def music():
    p=case('sozialversicherung-musikakademie-prenzlauer-berg','Musikakademie Prenzlauer Berg','Mara Noémi Zwirn unterrichtet seit 2022 Klavier an einer privaten Berliner Musikakademie. Im September 2026 treffen eine Betriebsprüfung, eine liegengebliebene Zustimmungserklärung und unterschiedliche Vorstellungen über die Unterrichtsorganisation aufeinander. Aktenstand: 30. September 2026.')
    doc(p,'01_Honorarvertrag_2022.docx','Honorarvertrag Klavierunterricht','Musikakademie am Mauerpark gGmbH · Kopenhagener Straße 74 · 10437 Berlin\nMara Noémi Zwirn · Schwedter Straße 188 · 10435 Berlin\nBerlin, 17. August 2022', '''1 Gegenstand und Laufzeit

Die Musikakademie am Mauerpark gGmbH, vertreten durch ihren Geschäftsführer Kunibert Schnurr, beauftragt Mara Noémi Zwirn ab dem 1. September 2022 mit der Erteilung von Einzelunterricht im Fach Klavier. Der Unterricht richtet sich an Kinder und Erwachsene, die mit der Akademie einen Unterrichtsvertrag schließen. Dieser Vertrag läuft auf unbestimmte Zeit. Der anfängliche Umfang beträgt voraussichtlich zwölf Unterrichtseinheiten pro Woche. Eine Unterrichtseinheit umfasst 45 Minuten; Vor- und Nachbereitung gehören zur vereinbarten Leistung.

2 Gestaltung des Unterrichts

Frau Zwirn entscheidet selbst über Methode, Literatur und Übefolgen. Die Akademie teilt ihr neue Unterrichtsinteressenten nach Rücksprache zu. Sie darf die Übernahme einzelner Schüler bei fehlender fachlicher Eignung oder unvereinbaren Terminen ablehnen. Vor Beginn eines Schulhalbjahres stimmen beide Seiten die Unterrichtstage ab. Nach Mitteilung der Termine an die Schüler sind Änderungen mit dem Büro abzustimmen, damit Räume nicht doppelt belegt werden. Das Büro führt den gemeinsamen Raumkalender.

Die Unterrichtsräume und die dort vorhandenen Klaviere dürfen für die vereinbarten Stunden kostenfrei genutzt werden. Außerhalb dieser Stunden bedarf die Nutzung der Zustimmung der Akademie. Ein Anspruch auf einen bestimmten Raum besteht nicht. Eigene Schüler darf Frau Zwirn in diesen Räumen nur aufgrund einer gesonderten Raumvereinbarung unterrichten. Die Akademie stellt weder einen eigenen Arbeitsplatz noch Bürozeiten für Frau Zwirn bereit.

3 Honorar und Ausfälle

Das Honorar beträgt 32 Euro je erteilter Unterrichtseinheit. Frau Zwirn legt bis zum fünften Werktag des Folgemonats eine Rechnung mit Datum, Schülerkürzel und Zahl der Einheiten vor. Die Akademie zahlt innerhalb von 14 Tagen nach Rechnungseingang. Für Schülerabsagen mit weniger als 24 Stunden Vorlauf wird das Honorar gezahlt. Bei rechtzeitigen Absagen sollen Ersatztermine angeboten werden; ohne Ersatzunterricht entsteht kein Honorar. Während der Berliner Schulferien findet grundsätzlich kein Regelunterricht statt. Gesondert vereinbarte Ferienkurse werden gesondert vergütet.

Fällt Frau Zwirn wegen Krankheit oder aus anderen persönlichen Gründen aus, besteht kein Honoraranspruch. Das Büro ist möglichst früh zu verständigen. Eine fachlich geeignete Vertretung kann nach vorheriger Mitteilung eingesetzt werden. Die Akademie kann widersprechen, wenn der Qualifikationsnachweis oder die für die Arbeit mit Minderjährigen erforderlichen Unterlagen fehlen. Vergütung und Abrechnung mit der Vertretung obliegen Frau Zwirn.

4 Veranstaltungen und Außendarstellung

Die Akademie veranstaltet jährlich ein Schülerkonzert. Eine Mitwirkung von Frau Zwirn wird jeweils gesondert vereinbart und mit 32 Euro je 45 Minuten vergütet. Eine Verpflichtung zur Teilnahme an Teamtreffen besteht nicht. Die Akademie darf Namen, Kurzvita und ein von Frau Zwirn freigegebenes Bild für die Information ihrer Schüler verwenden. Die Kontaktdaten der Schüler werden ausschließlich für Unterrichtsorganisation und Kommunikation bereitgestellt.

5 Weitere Tätigkeiten und Beendigung

Frau Zwirn darf für andere Auftraggeber arbeiten und eigene Unterrichtsangebote betreiben. Aktuelle Akademieschüler darf sie während des laufenden Unterrichtsvertrags nicht unter Nutzung der überlassenen Kontaktdaten für eigene Angebote abwerben. Beide Seiten können zum Ende eines Kalendermonats mit sechs Wochen Frist schriftlich kündigen. Bereits abgerechnete Leistungen werden nach Vertragsende innerhalb von 14 Tagen bezahlt. Schlüssel und überlassene Schülerlisten sind zurückzugeben.

6 Vertragsverständnis

Beide Seiten gehen von einer selbständigen Honorartätigkeit aus. Frau Zwirn führt ihre Steuern selbst ab und kümmert sich um ihre Krankenversicherung und Altersvorsorge. Ein Urlaubsanspruch wird nicht vereinbart. Änderungen dieses Vertrags sollen in Textform festgehalten werden. Mündliche Absprachen über einzelne Unterrichtstermine bleiben möglich.

Berlin, 17. August 2022

Kunibert Schnurr, Geschäftsführer · Mara Noémi Zwirn

Ausfertigung für Frau Zwirn. Die beiderseits unterzeichnete Papierfassung befindet sich im Personalordner des Büros.''','Kunibert Schnurr')
    doc(p,'02_Nachtrag_Honorar_2024.docx','Nachtrag zum Honorarvertrag','Musikakademie am Mauerpark gGmbH / Mara Noémi Zwirn\nBerlin, 8. Januar 2024', '''1 Honorar ab Januar 2024

Für Unterrichtseinheiten, die ab dem 1. Januar 2024 erbracht werden, erhöht sich das Honorar aus Ziffer 3 des Vertrags vom 17. August 2022 auf 36 Euro je 45 Minuten. Die Erhöhung gilt ebenfalls für entschädigte kurzfristige Schülerabsagen und vereinbarte Mitwirkung beim Schülerkonzert. Bereits im Dezember 2023 erteilte Stunden bleiben bei 32 Euro, auch wenn die Rechnung erst im Januar eingeht.

2 Vorbereitung der Studienbewerbung

Frau Zwirn übernimmt zusätzlich die Betreuung von bis zu vier Schülerinnen und Schülern im Kurs „Mappe Musik“. Der Kurs bereitet auf Aufnahmeprüfungen an Musikhochschulen vor; die Akademie ist selbst keine Hochschule und verleiht keine staatlichen Abschlüsse. Über Repertoire und Prüfungsprogramm entscheidet Frau Zwirn im Gespräch mit den Teilnehmenden. Das Büro führt die Anmeldungen und stimmt die Probespieltermine mit dem Gehörbildungskurs ab.

Im Schulhalbjahr Februar bis Juli 2024 stellt Frau Zwirn für diese vier Teilnehmenden jeweils eine kurze Einschätzung für das Abschlussgespräch bereit. Ein von der Akademie bereitgestelltes Formular kann verwendet werden. Ein gesondertes Honorar für die Einschätzung wird nicht gezahlt. Das Abschlussgespräch selbst wird nach tatsächlicher Dauer in Unterrichtseinheiten vergütet.

3 Organisatorische Abstimmung

Die Akademie bittet darum, Fehlzeiten und Änderungen von Stunden jeweils bis zum folgenden Montag im Portal einzutragen. Diese Einträge ermöglichen die Abrechnung gegenüber den Familien. Die Beteiligten sind sich darüber einig, dass der Vertrag keine allgemeine Anwesenheits- oder Bereitschaftspflicht begründet. Frau Zwirn soll bei geplanter längerer Abwesenheit vier Wochen vorher mitteilen, welche Ersatztermine sie anbieten kann. Die Regelung zur Vertretung in Ziffer 3 bleibt bestehen.

4 Fortgeltung

Die übrigen Regelungen des Vertrags vom 17. August 2022 gelten unverändert. Dieser Nachtrag verlängert keine Kündigungsfrist und begründet keinen Mindestumfang. Die anfänglich erwarteten zwölf Einheiten sind durch zusätzliche Schüler inzwischen überschritten; die Parteien treffen hierzu keine garantierte Abnahmevereinbarung.

Berlin, 8. Januar 2024

Kunibert Schnurr · Mara Noémi Zwirn

Dateinotiz Büro: Beide Papierunterschriften liegen seit dem 12. Januar 2024 vor. Die Ausfertigung wurde Frau Zwirn bei Abholung des Unterrichtsschlüssels ausgehändigt.''','Kunibert Schnurr')
    doc(p,'03_Unterrichtsbetrieb_September_2025.docx','Absprachen für das Winterhalbjahr','Büro der Musikakademie am Mauerpark\nArbeitsstand vom 28. August 2025 · versandt an alle Lehrenden', '''1 Stunden und Räume

Ab dem 8. September verwenden wir ausschließlich den Kalender im Portal. Die Raumbelegung aus den E-Mails vom Juni wird dadurch ersetzt. Mara hat dienstags Raum 3 von 14 bis 19 Uhr und donnerstags Raum 2 von 14 bis 20 Uhr. Bei Stimmarbeiten kann das Büro auf Raum 4 umbuchen. Bitte verschiebt keine Unterrichtsstunde ohne Eintrag: Frau Becker und Herr Lenz hatten letzte Woche denselben Raum, obwohl beide Familien eine Bestätigung bekommen hatten.

Für den September sind neben diesen Regelzeiten zusätzliche Räume für Maras Einzelunterricht und Nachholtermine reserviert: mittwochs am 10., 17. und 24. September jeweils Raum 4 von 10 bis 15.15 Uhr (sieben Einheiten à 45 Minuten), freitags am 12. und 26. September jeweils Raum 3 von 10 bis 14.30 Uhr (sechs Einheiten). Die zusätzlichen 33 Zeitfenster gelten nur für September; zwei der insgesamt 81 Septemberfenster bleiben zunächst für Terminwechsel frei. Montags hat Mara den anderen Auftrag in Pankow. Hannes trägt die bestätigten Familien einzeln ein. Abgerechnet werden die tatsächlich erteilten oder nach dem Vertrag vergüteten Einheiten, nicht die Raumreservierungen.

Die Einträge sind zunächst unser Planungsvorschlag. Rückmeldungen bitte bis zum 3. September an Hannes. Nach den Elternbestätigungen brauchen wir verlässliche Termine. Wer eine Lücke zwischen zwei Stunden anders nutzen möchte, muss selbstverständlich nicht im Haus bleiben. Die Eingangstür schließt ab 18 Uhr; der Schlüssel darf nicht an Schüler weitergegeben werden.

2 Abwesenheit und Vertretung

Im Portal heißt die Funktion technisch „Urlaubsantrag“. Auch Honorarkräfte tragen dort ihre geplanten Abwesenheiten ein. Das Büro bestätigt anschließend, ob die betroffenen Familien informiert sind. In der letzten Sitzung wurde ausdrücklich gefragt, ob das eine Genehmigung sei. Kunibert möchte die Frage mit der Verwaltung klären; bis dahin bleibt die Bezeichnung im System unverändert. Kurzfristige Erkrankungen bitte per Telefon mitteilen, weil wir die Portalbenachrichtigung nicht ständig überwachen.

Wer eine Vertretung vorschlägt, sendet Name, Telefonnummer und Qualifikationsnachweis ans Büro. Bei einer Wiederholungsvertretung müssen die Unterlagen nicht erneut eingereicht werden. Die Zahlung klärt die vertretene Lehrkraft direkt. Hannes vermerkt im Portal den Namen der tatsächlich unterrichtenden Person.

3 Treffen und Veranstaltungen

Die Besprechung zur Probespielwoche findet am 15. September um 10 Uhr statt. Für Lehrkräfte im Kurs „Mappe Musik“ brauchen wir die Rückmeldungen zu allen Teilnehmern spätestens an diesem Tag. Eine schriftliche Rückmeldung reicht, falls der Termin nicht passt. Das Büro hat in der Rundmail versehentlich „Teilnahme verpflichtend“ geschrieben; die Korrektur ist noch nicht an alle Verteiler gegangen.

Der Konzerttermin am 13. Dezember steht fest. Die Mitwirkung wird wie im Vorjahr mit zwei Einheiten abgerechnet. Wer nicht mitwirkt, soll die betroffenen Familien frühzeitig informieren. Bitte sagt Konzertbeiträge nicht gegenüber Eltern verbindlich zu, bevor wir die gesamte Programmlänge geprüft haben.

Hannes Riedel, Unterrichtsorganisation''','Hannes Riedel')
    doc(p,'04_Erinnerung_Mara_20260924.docx','Mein Unterricht an der Akademie','Mara Noémi Zwirn\nBerlin, 24. September 2026 · für Rechtsanwältin Vera Sommer', '''Ich möchte, dass Sie meine Situation prüfen, bevor ich irgendetwas unterschreibe. Ich habe im Büro immer gesagt, dass ich die Arbeit gern behalten möchte. Das war keine Erklärung, auf eine Absicherung zu verzichten. Die Datei mit der Zustimmung habe ich 2025 geöffnet und ausgefüllt, aber nach meiner Erinnerung nicht abgeschickt. Auf meinem alten Laptop steht allerdings ein PDF im Ordner „gesendet“. Ich weiß nicht mehr, ob Hannes es im Büro vom USB-Stick kopiert hat.

Zu Beginn hatte ich tatsächlich nur zwölf Einheiten. Ich konnte die beiden Tage passend zu meinen Konzerten legen. Seit Herbst 2024 kamen immer mehr Termine hinzu. Meist habe ich neue Schüler angenommen, zweimal habe ich abgelehnt. Bei einer Schülerin war es wegen eines schwierigen Elterngesprächs; die andere wollte montags kommen. Beide E-Mails müssten im Postfach sein. Eigene Preise habe ich mit den Familien nie besprochen. Die Akademie wollte ausdrücklich, dass alle Zahlungsfragen über das Büro laufen.

Ich habe mich 2023 einmal vertreten lassen. Lea bekam das Geld von mir. Im Portal stand trotzdem zunächst mein Name. Beim zweiten Versuch im November 2025 sagte Hannes, Leas erweitertes Führungszeugnis sei zu alt. Ich habe dann die Stunden selbst nachgeholt. Wenn ich krank war, wurde kein Honorar gezahlt. Für Schüler, die am selben Tag abgesagt haben, bekam ich regelmäßig Geld. Einmal strich die Buchhaltung vier Einheiten, weil Hannes die Absagen als rechtzeitig eingetragen hatte. Darum stimmt mein Januarzettel nicht mit der endgültigen Rechnung überein.

Ich unterrichte außerdem privat, aber zu Hause nur samstags. Drei dieser Schüler kenne ich von vor der Akademiezeit. Es gibt keine Angestellten. Mein Mann hilft gelegentlich beim Tragen des E-Pianos für ein Konzert; dafür bezahlt ihn niemand. 2025 habe ich nach meiner Liste 8.460 Euro aus Privatunterricht und Auftritten bekommen. Bei einem Auftritt waren Fahrtkosten enthalten. Die Aufstellung wurde nicht für die Steuer gemacht und enthält noch keine Ausgaben.

Die Künstlersozialkasse hat mir im Mai 2026 Fragen geschickt. Ich habe auf die Nachfrage noch nicht vollständig geantwortet, weil ich die Trennung zwischen Unterricht und Konzert nicht sicher hinbekam. Ich war bis jetzt freiwillig bei einer gesetzlichen Krankenkasse versichert und habe zusätzlich einen privaten Rentensparplan mit 150 Euro im Monat. Ich dachte, das reiche. Einen Bescheid der Rentenversicherung zur selbständigen Lehrtätigkeit habe ich nicht gefunden.

Bitte schreiben Sie der Akademie zunächst nicht. Ich brauche am 6. Oktober eine Einschätzung für das Gespräch mit Kunibert. Hannes sagte am Telefon, die Prüfung betreffe nur die Einrichtung. Im Schreiben werden aber ausdrücklich meine Unterlagen angefordert. Das verunsichert mich. Mir ist wichtig, dass der Unterricht nicht einfach ausfällt und meine Schüler nicht in den Streit geraten.''','Mara Noémi Zwirn')
    doc(p,'05_Protokoll_Buero_20260925.docx','Gespräch über die Unterlagenanforderung','Musikakademie am Mauerpark gGmbH\n25. September 2026, 11.10 bis 11.45 Uhr\nTeilnehmende: Kunibert Schnurr, Hannes Riedel, Nils Brandt (Buchhaltung)', '''Kunibert bittet darum, die angeforderten Honorardaten unverändert aus der Buchhaltung zu exportieren. Eine Umbenennung früherer Dateien soll unterbleiben. Nils hat die Jahre 2022 bis 2025 bereits zusammengestellt. Für Januar 2025 liegen zwei Rechnungsstände vor; nur die korrigierte Rechnung mit 2.448 Euro wurde gebucht. Die zuerst eingereichten 2.592 Euro wurden nicht bezahlt. Die Differenz betrifft vier abgesagte Unterrichtseinheiten zu je 36 Euro.

Hannes erinnert sich, dass Mara die Zustimmung im Mai 2025 im Büro abgegeben habe. Auf Nachfrage kann er weder Datum noch Übermittlungsweg nennen. Im Portal steht beim Feld „Erklärung erhalten“ der 19. Mai 2025. Dieser Eintrag wurde bei der Sammelübernahme von Hannes erzeugt. Die Datei im Ordner von Mara trägt den Dateinamen „Zustimmung_Mara_final.pdf“, aber keine erkennbare Unterschrift. Das Datum 19. Mai steht im Textfeld. Nils weist darauf hin, dass der Import auch die Vorlagen aus dem gemeinsamen Ordner mitgenommen habe.

Kunibert möchte Mara um eine erneute Erklärung bitten. Hannes soll dabei nicht behaupten, dass eine frühere Unterschrift bewiesen sei. Mara hatte zuletzt geäußert, sie wolle vor einer Entscheidung beraten werden. Die Unterrichtsplanung für Oktober wird unabhängig davon fortgeführt. Über eine Anstellung wurde noch kein Angebot gemacht; Kunibert hat lediglich im Frühjahr überschlägige Personalkosten rechnen lassen. Eine Stellenbeschreibung existiert nicht.

Zum Portal hält Hannes fest, dass „Urlaubsantrag“ ein Standardfeld des Anbieters ist. Bei Honorarkräften sei die Freigabe aus seiner Sicht nur eine Kommunikationskontrolle. Kunibert erinnert jedoch an zwei Fälle, in denen eine Reise wegen des Probespiels verschoben wurde. Ob dies auf Bitte oder Anweisung geschah, lässt sich im Gespräch nicht klären. Die ursprünglichen Nachrichten sollen gesucht werden.

Nils wird die Aufstellung um Zahlungsdatum und Korrekturbeleg ergänzen. In der Auswertung werden Unterrichtseinheiten und bezahlte Ausfälle getrennt dargestellt. Er möchte außerdem wissen, ob Zahlungen für Konzerte in der Honorarliste enthalten sein sollen. Kunibert sagt, sie sollen vollständig und mit ihrem ursprünglichen Buchungstext erscheinen. Die Einordnung werde nicht von der Buchhaltung entschieden.

Protokolliert von Nils Brandt am 25. September 2026. Hannes bestätigte am selben Nachmittag nur seine Angaben zum Datenimport; Kunibert hat den Text noch nicht gegengelesen.''','Nils Brandt')
    pdf(p,'06_Pruefankuendigung_20260910.pdf','Prüfung der Arbeitgeberunterlagen','Deutsche Rentenversicherung Bund · Prüfdienst\nAn Musikakademie am Mauerpark gGmbH\nBerlin, 10. September 2026 · Geschäftszeichen 26-BA-4817', '''Sehr geehrter Herr Schnurr,

für Ihren Betrieb ist eine Prüfung nach Paragraf 28p SGB IV vorgesehen. Der Prüfzeitraum umfasst den 1. Januar 2022 bis zum 31. Dezember 2025. Als Beginn der Prüfung ist der 12. Oktober 2026 vorgesehen. Bitte stellen Sie die elektronischen Entgeltunterlagen und die nachstehend bezeichneten ergänzenden Unterlagen bis zum 2. Oktober 2026 bereit. Falls einzelne Unterlagen nicht innerhalb dieser Zeit verfügbar sind, teilen Sie dies bitte unter Angabe des voraussichtlichen Bereitstellungstermins mit.

Benötigt werden die Sachkonten für Honorare, Fremdleistungen und sonstige Unterrichtsvergütungen, die zugrunde liegenden Verträge einschließlich Änderungen sowie eine Zuordnung der Zahlungen zu den eingesetzten Personen. Für Frau Mara Noémi Zwirn bitten wir zusätzlich um Stunden- und Raumpläne, Abrechnungen ausgefallener Stunden, vorhandene Vertretungsvereinbarungen und Unterlagen über die tatsächliche Unterrichtsorganisation. Teilen Sie bitte mit, ob zu dieser Tätigkeit bereits ein Statusfeststellungsverfahren oder ein Verfahren bei einer Einzugsstelle geführt wurde.

Soweit Sie sich hinsichtlich der Lehrtätigkeiten auf Paragraf 127 SGB IV berufen, übersenden Sie bitte die hierzu vorhandenen Erklärungen und Nachweise ihres Zugangs. Bitte reichen Sie vorhandene Fassungen mit abweichenden Angaben vollständig ein und erläutern Sie deren Entstehung. Mit dieser Anforderung ist noch keine Feststellung zum Erwerbsstatus oder zur Beitragspflicht einzelner Personen verbunden.

Die Prüfung erfolgt zunächst anhand der übermittelten Unterlagen. Eine persönliche Besprechung wird erforderlichenfalls gesondert abgestimmt. Bitte benennen Sie eine Person, die Rückfragen zur Buchhaltung und zum Unterrichtsbetrieb beantworten kann.

Mit freundlichen Grüßen
Im Auftrag
Sabine Hertel''')
    pdf(p,'07_Rechnung_Januar_2025_Erstfassung.pdf','Rechnung MS 2025 01','Mara Noémi Zwirn · Schwedter Straße 188 · 10435 Berlin\nAn Musikakademie am Mauerpark gGmbH\nRechnungsdatum 3. Februar 2025', '''Ich berechne den im Januar 2025 durchgeführten Klavierunterricht und die nach dem Honorarvertrag vergüteten kurzfristigen Schülerabsagen. Meine Liste umfasst 64 tatsächlich erteilte Einheiten und acht kurzfristige Absagen. Die in dieser Rechnung enthaltenen Absagen vom 9. und 16. Januar wurden mir nach meiner Erinnerung jeweils am Unterrichtstag mitgeteilt. Zu den Ersatzterminen im Februar führe ich eine eigene Liste; diese Stunden sind hier nicht enthalten.

Es ergeben sich 72 Einheiten zu jeweils 36 Euro, insgesamt 2.592 Euro. Umsatzsteuer wird nicht ausgewiesen. Die Rechnung soll innerhalb von 14 Tagen überwiesen werden. Als Zahlungsreferenz bitte ausschließlich MS 2025 01 verwenden. Die Buchhaltung besitzt bereits meine unveränderte Bankverbindung.

Die Stundenaufstellung wurde am 2. Februar aus meinem Kalender übertragen. Das Schülerkürzel HB erscheint zweimal am 23. Januar, weil an diesem Tag eine Doppelstunde vereinbart war. Für die Schülerin MK wurde keine Januarstunde berechnet; sie war während des gesamten Monats abgemeldet. Bitte melden Sie sich, falls die Portalangaben von meinem Kalender abweichen.

Mara Noémi Zwirn''',table=[['Leistung','Einheiten','Satz EUR','Betrag EUR'],['Unterricht Januar',64,'36,00','2.304,00'],['Kurzfristige Absagen',8,'36,00','288,00'],['Gesamt',72,'','2.592,00']])
    pdf(p,'08_Rechnung_Januar_2025_korrigiert.pdf','Korrigierte Rechnung MS 2025 01 K','Mara Noémi Zwirn · Schwedter Straße 188 · 10435 Berlin\nAn Musikakademie am Mauerpark gGmbH\nRechnungsdatum 10. Februar 2025', '''Diese Rechnung ersetzt meine Rechnung MS 2025 01 vom 3. Februar 2025 vollständig. Bitte zahlen Sie ausschließlich auf diese korrigierte Rechnung. Die Zahl der tatsächlich erteilten Einheiten bleibt bei 64. Nach Rücksprache mit Hannes werden von den acht zunächst eingetragenen kurzfristigen Absagen nur vier im Januar bezahlt. Damit ergeben sich 68 Einheiten zu jeweils 36 Euro und ein Gesamtbetrag von 2.448 Euro.

Die vier gestrichenen Einheiten betreffen die Portalvorgänge A-119, A-120, A-137 und A-138. Ich hatte sie auf meinem Kalender als kurzfristig abgesagt markiert. Hannes hat mir am 7. Februar erklärt, dass die Nachrichten der Eltern bereits am Vortag im Sammelpostfach eingegangen seien. Ich habe die Kürzung vorläufig übernommen, damit die Zahlung nicht weiter liegenbleibt. Wir wollen die Ersatztermine gesondert klären.

Für die Zahlung gilt die bereits hinterlegte Bankverbindung. Die Vergütung enthält die Vorbereitung des Unterrichts; Fahrtkosten stelle ich nicht zusätzlich in Rechnung. Umsatzsteuer wird nicht ausgewiesen. Die Zahlungsfrist läuft bis zum 24. Februar 2025.

Mara Noémi Zwirn''',table=[['Leistung','Einheiten','Satz EUR','Betrag EUR'],['Unterricht Januar',64,'36,00','2.304,00'],['Vergütete Absagen',4,'36,00','144,00'],['Gesamt',68,'','2.448,00']])
    pdf(p,'09_Zahlungsbeleg_Februar_2025.pdf','Kontoumsätze Februar 2025','Export aus dem Geschäftskonto der Musikakademie am Mauerpark gGmbH\nErstellt von Nils Brandt am 23. September 2026', '''Der Auszug enthält die drei zum Suchbegriff „Zwirn“ gefundenen Buchungen vom 1. bis zum 28. Februar 2025. Die Auswertung wurde aus der Buchhaltungsansicht erstellt. Sie enthält keine weiteren Kontoumsätze und weist daher keinen vollständigen Kontoanfangs- oder Endbestand aus.

Am 18. Februar wurde die korrigierte Januarrechnung MS 2025 01 K in Höhe von 2.448 Euro überwiesen. Am selben Tag wurde eine separate Erstattung für gemeinsam beschaffte Noten in Höhe von 27,80 Euro ausgeführt. Die am 4. Februar angelegte Zahlung über 2.592 Euro wurde vor Freigabe gelöscht und hat keinen Bankumsatz ausgelöst. Sie erscheint deshalb nicht in der unten stehenden Tabelle. Der vorbereitete Zahlungslauf liegt in einem anderen Archiv.

Die weitere Zahlung über 72 Euro vom 26. Februar gehört zum Probespiel am 15. Februar und nicht zum Januarunterricht. Der Buchungstext verwendet versehentlich „Konzert Januar“. Frau Zwirn hatte diese Leistung separat als zwei Einheiten abgerechnet. Der Betrag ist in der Honorarauswertung 2025 im Monat Februar enthalten.

Nils Brandt, Buchhaltung''',table=[['Valuta','Empfänger','Verwendungszweck','Soll EUR'],['18.02.2025','Mara Noémi Zwirn','MS 2025 01 K','2.448,00'],['18.02.2025','Mara Noémi Zwirn','Auslage Noten NB-18','27,80'],['26.02.2025','Mara Noémi Zwirn','Konzert Januar MS-02-P','72,00']])
    pdf(p,'10_Portal_Abwesenheit_202506.pdf','Abwesenheitseintrag 447','Musikakademie am Mauerpark · Portalexport\nPerson Mara Noémi Zwirn · Export 22. September 2026', '''Erfasst am 4. Juni 2025 um 21.14 Uhr von Mara Noémi Zwirn: Abwesenheit vom 23. bis zum 27. Juni 2025, Grund „Konzertreise“. Betroffen sind die Unterrichtstage Dienstag und Donnerstag. Als Ersatz sind Samstag, 14. Juni, und Samstag, 5. Juli, genannt. Der Eintrag wurde durch das Portal als Urlaubsantrag geführt.

Kommentar von Hannes Riedel am 5. Juni um 09.02 Uhr: „Bitte noch nicht den Eltern zusagen. Der Samstag im Juli ist wegen des Sommerfestes dicht. Für die vier Mappenschüler brauchen wir eine andere Lösung. Kannst du die Reise um einen Tag verschieben?“

Kommentar von Mara Noémi Zwirn am 5. Juni um 11.31 Uhr: „Die Reise ist gebucht. Ich kann Donnerstagabend nach Rückkehr zwei Doppelstunden machen, aber nur wenn Raum 2 frei ist. Die übrigen Eltern frage ich nach Sonntag. Das ist diesmal wirklich aufwendig.“

Am 9. Juni 2025 um 15.46 Uhr änderte Hannes Riedel den Status auf „genehmigt“. Der automatisch erzeugte Text lautete „Ihr Urlaubsantrag wurde freigegeben“. Die Benachrichtigung ging an mara.zwirn@klavierraum.example. Die Historie enthält keine Änderung des Abwesenheitszeitraums. Die ursprüngliche Elterninformation ist in diesem Export nicht enthalten. Der Portalbetreiber speichert Uhrzeiten in der Zeitzone Europe/Berlin.''')
    pdf(p,'11_Vertretung_Lea_202310.pdf','Abrechnung der Unterrichtsvertretung','Lea Winter · Stargarder Straße 82 · 10437 Berlin\nAn Mara Noémi Zwirn · Rechnung LW 2023 17 · 31. Oktober 2023', '''Für die von mir am 10. und 12. Oktober 2023 übernommene Vertretung berechne ich zehn Klavierunterrichtseinheiten zu je 30 Euro, insgesamt 300 Euro. Die Einheiten dauerten jeweils 45 Minuten. Ich habe die Schüler in den Räumen der Musikakademie am Mauerpark unterrichtet und die Anwesenheit auf der mir überlassenen Papierliste vermerkt. Mara Noémi Zwirn hatte mir vorab die Stücke und den jeweiligen Arbeitsstand geschickt.

Die Abstimmung über Termine und Vergütung erfolgte zwischen Mara Noémi Zwirn und mir. Hannes Riedel bestätigte mir am 9. Oktober telefonisch, dass die hinterlegten Unterlagen aus meinem Ferienkurs vom Sommer weiterhin ausreichen. Einen eigenen Unterrichtsvertrag mit der Akademie habe ich für diese Vertretung nicht geschlossen. Den Schlüssel holte ich bei Mara ab und gab ihn am 13. Oktober zurück.

Bitte überweisen Sie den Gesamtbetrag bis zum 14. November 2023 auf meine bekannte Bankverbindung. Fahrtkosten werden nicht gesondert berechnet. Umsatzsteuer wird nicht ausgewiesen. Zwei zunächst vorgesehene Stunden am 12. Oktober wurden von den Familien rechtzeitig abgesagt und sind nicht berechnet. Die beigefügte Stundenliste umfasst deshalb zehn und nicht zwölf Einheiten.

Lea Winter''')
    pdf(p,'12_Zustimmung_Mara_final.pdf','Erklärung zur Lehrtätigkeit','Musikakademie am Mauerpark gGmbH\nDateistand im Personalordner: 19. Mai 2025', '''Name der Lehrkraft: Mara Noémi Zwirn. Vertragspartner: Musikakademie am Mauerpark gGmbH. Tätigkeit: Klavierunterricht auf Grundlage des Vertrags vom 17. August 2022 und des Nachtrags vom 8. Januar 2024.

Die Vertragsparteien sind bei Abschluss des Vertrags davon ausgegangen, dass die Lehrtätigkeit selbständig ausgeübt wird. Ich stimme zu, dass die Versicherungspflicht aufgrund einer Beschäftigung für die bezeichnete Lehrtätigkeit im Rahmen der gesetzlichen Übergangsregelung erst nach Ablauf der Übergangszeit eintritt. Die Erklärung bezieht sich ausschließlich auf die Tätigkeit für den genannten Vertragspartner.

Ich habe die Information erhalten, dass meine bisherige Krankenversicherung fortzuführen ist und dass Fragen meiner eigenen Rentenversicherung beziehungsweise einer Versicherung über die Künstlersozialkasse gesondert zu klären sind. Ich möchte vor Abgabe der Erklärung noch eine Rückmeldung dazu, ob meine bisherigen privaten Sparbeiträge berücksichtigt werden. Die Ansprechpartnerin im Büro konnte mir dies nicht beantworten.

Im Datumsfeld ist „19.05.2025“ eingetragen. Das Unterschriftsfeld in der gespeicherten Datei enthält keine Unterschrift. Eine Übermittlungsbestätigung ist der Datei nicht beigefügt. Der letzte Absatz wurde von Mara Noémi Zwirn in das ursprünglich vorgesehene Freitextfeld geschrieben. Die Datei wurde vom Büro am 23. September 2026 unverändert aus dem gemeinsamen Personalordner exportiert.''')
    pdf(p,'13_KSK_Nachfrage_20260518.pdf','Ihre Angaben zur künstlerischen Tätigkeit','Künstlersozialkasse\nAn Frau Mara Noémi Zwirn · Berlin\n18. Mai 2026 · Vorgang K-26-31842', '''Sehr geehrte Frau Zwirn,

vielen Dank für Ihre Unterlagen vom 27. April 2026. Für die Bearbeitung Ihres Antrags benötigen wir ergänzende Angaben zur Art und zum Umfang Ihrer Tätigkeiten. Bitte unterscheiden Sie den Klavierunterricht für die Musikakademie am Mauerpark, Ihren eigenen Privatunterricht und Ihre Konzerttätigkeit. Erläutern Sie, wie sich die Honorare im laufenden Jahr voraussichtlich auf diese Bereiche verteilen und welche Betriebsausgaben dabei entstehen.

Bitte übersenden Sie den Vertrag mit der Akademie einschließlich Nachträgen sowie zwei aktuelle Rechnungen. Falls ein Versicherungsträger bereits eine Entscheidung zum Erwerbsstatus dieser Tätigkeit getroffen hat, benötigen wir eine Kopie des vollständigen Bescheids. Sofern eine Erklärung zur Übergangsregelung für Lehrtätigkeiten abgegeben wurde, fügen Sie diese bitte ebenfalls bei und teilen Sie das Datum der Abgabe mit.

Aus den eingereichten Kontoauszügen geht bisher nicht hervor, ob Sie Personen für Ihren Unterrichtsbetrieb beschäftigen. Bitte beantworten Sie diese Frage ausdrücklich. Die unentgeltliche Hilfe eines Familienangehörigen bei einzelnen Veranstaltungen beschreiben Sie bitte getrennt. Ihre Beitragszahlungen an eine freiwillige Krankenversicherung sind durch die übersandten Auszüge belegt; eine Entscheidung über Ihren Antrag ist damit noch nicht getroffen.

Wir bitten um Antwort bis zum 15. Juni 2026. Bei Rückfragen geben Sie bitte das oben genannte Vorgangszeichen an.

Mit freundlichen Grüßen
Im Auftrag
Kerstin Lübke''')
    pdf(p,'14_Portalimport_20250519.pdf','Importlauf Personalunterlagen','Büroprotokoll Datenübernahme\nLauf 2025-05-19-02 · 19. Mai 2025, 18.06 Uhr', '''Der Import wurde durch Benutzer h.riedel aus dem Verzeichnis Teamablage/Lehrkraefte/Zustimmung gestartet. Das Programm ordnete Dateien anhand der ersten übereinstimmenden Namensfolge einem Personalprofil zu. Die Auswahl „bestehende Datumsfelder erhalten“ war nicht aktiviert. Für jedes zugeordnete Dokument wurde das Feld „Erklärung erhalten“ auf das Datum des Importlaufs gesetzt.

Es wurden 17 Dateien geprüft, 15 automatisch zugeordnet und zwei wegen doppelter Namensbestandteile zurückgestellt. Die Datei Zustimmung_Mara_final.pdf wurde dem Profil Mara Noémi Zwirn zugeordnet. Die Verarbeitung prüfte Dateiformat und Lesbarkeit. Eine Prüfung auf eine Unterschrift oder auf den Inhalt des Freitextfelds fand nicht statt. Der Quellordner enthielt zum Zeitpunkt des Imports sowohl ausgefüllte Formulare als auch bearbeitete Entwürfe.

Hannes Riedel ergänzte am 20. Mai den Vermerk: „Rücklauf bei den beiden offenen Namen manuell kontrollieren; die übrigen Einträge sehen vollständig aus.“ Dieser Vermerk wurde nicht um eine Liste einzelner Zugangsbelege ergänzt. Ein Abgleich mit dem E-Mail-Postfach ist im Laufprotokoll nicht verzeichnet. Das Programm löschte die Quelldateien nicht.

Exportiert am 23. September 2026 durch Nils Brandt. Die Liste zeigt den technischen Importvorgang; sie enthält keine Angaben dazu, wer die jeweilige Datei ursprünglich in der Teamablage gespeichert hat.''')
    mails=[
    ('15_Auftrag_an_Kanzlei.eml','2026-09-23T19:24:00+02:00','Mara Noémi Zwirn <mara.zwirn@klavierraum.example>','Vera Sommer <vs@sommer-kanzlei.example>','Unterricht Akademie und Schreiben der Rentenversicherung', '''Sehr geehrte Frau Sommer,

ich schicke Ihnen das Schreiben, das mir Hannes heute weitergeleitet hat. Ich unterrichte seit vier Jahren in der Akademie. Eigentlich wollte ich damals bewusst flexibel bleiben. Jetzt sagt Kunibert, ich müsse nur noch eine Erklärung nachreichen und alles könne so weitergehen. Gleichzeitig sollen wir unsere Zeiten sehr genau nachweisen.

Bitte prüfen Sie für mich, was meine Unterschrift bewirken würde und ob ich für die vergangenen Jahre noch selbst etwas zahlen muss. Ich möchte nicht sofort kündigen. Ein festes Arbeitsverhältnis wäre für mich inzwischen interessant, aber ich möchte auch meine Konzertreisen behalten. Ich habe der Akademie noch keine Vollmacht oder Freigabe für Gespräche mit Ihnen gegeben.

Ich bringe morgen den alten Laptop vorbei. Dort liegen zwei Fassungen der Erklärung und meine Jahreslisten. Den Brief der Künstlersozialkasse habe ich bisher nicht beantwortet; er liegt ebenfalls bei.

Mit freundlichen Grüßen
Mara Noémi Zwirn
Schwedter Straße 188, 10435 Berlin''',None,('06_Pruefankuendigung_20260910.pdf','13_KSK_Nachfrage_20260518.pdf')),
    ('16_Raumplanung_2024.eml','2024-08-26T10:16:00+02:00','Hannes Riedel <buero@mauerpark-musik.example>','Mara Noémi Zwirn <mara.zwirn@klavierraum.example>','Neue Schüler und donnerstags Raum 2', '''Hallo Mara,

wir haben sechs neue Anfragen für Donnerstag. Ich habe den Tag zunächst von 14 bis 20 Uhr für dich geblockt. Bitte sag bis Freitag, welche Schüler du übernehmen kannst. Bei den vier Leuten für die Studienvorbereitung brauche ich nach Möglichkeit eine Zusage für das ganze Halbjahr, weil Gehörbildung und Probespiel darauf aufbauen.

Deinen Wunsch nach einer Pause um 16 Uhr konnte ich nur an zwei Wochen erfüllen. Sonst würde eine Familie kündigen, die schon seit 2021 hier ist. Wenn du mit dem späteren Ende nicht einverstanden bist, müssen wir die zwei letzten Termine anderweitig besetzen. Trage die Entscheidung bitte nicht nur in deinen eigenen Kalender ein, sondern antworte mir kurz.

Viele Grüße
Hannes
Musikakademie am Mauerpark, Unterrichtsbüro''',None,()),
    ('17_Ablehnung_Schueler.eml','2024-08-27T08:39:00+02:00','Mara Noémi Zwirn <mara.zwirn@klavierraum.example>','Hannes Riedel <buero@mauerpark-musik.example>','Re: Neue Schüler und donnerstags Raum 2', '''Hallo Hannes,

die beiden letzten Termine kann ich übernehmen. Die Pause um 16 Uhr brauche ich aber wegen der Betreuung meines Sohnes wenigstens jeden zweiten Donnerstag. Bitte lege da nicht wieder spontan eine Probestunde hinein. Von den sechs Anfragen lehne ich Herrn P. ab. Das passt fachlich nicht zu meinem Schwerpunkt; er sucht Jazzimprovisation. Die fünf anderen würde ich nach den Probestunden gern unterrichten.

Am Montag werde ich auch dieses Halbjahr nicht kommen. Dass das Büro dort noch eine Lücke hat, weiß ich, aber ich habe dafür bereits die andere Schule zugesagt. Für die Studienbewerber kann ich die vier Termine donnerstags halten. Das Repertoire spreche ich direkt mit ihnen ab; ich verwende dafür kein gemeinsames Kursbuch.

Mara''','16_Raumplanung_2024.eml',()),
    ('18_Zustimmung_Nachfrage_2025.eml','2025-05-21T22:05:00+02:00','Mara Noémi Zwirn <mara.zwirn@klavierraum.example>','Hannes Riedel <buero@mauerpark-musik.example>','Formular noch nicht freigegeben', '''Hallo Hannes,

ich habe gesehen, dass im Portal jetzt bei allen ein Datum steht. Meine Datei ist noch nicht zur Verwendung freigegeben. Ich habe sie am Montag ausgefüllt und die Frage zur Rentenversicherung hineingeschrieben. Bitte kläre erst, ob mein privater Vertrag reicht oder ob ich zusätzlich zur Rentenversicherung muss. Ich möchte keine Erklärung abgeben, deren Folgen ich nicht verstehe.

Falls du das PDF vom Stick schon abgelegt hast: Bitte lass es vorerst bei der Korrespondenz und nicht bei den unterschriebenen Verträgen. Ich hatte dir den Stick gegeben, weil darauf auch die Konzertprogramme waren. Du sagtest, das Formular müsse nicht unbedingt handschriftlich unterschrieben werden. Trotzdem möchte ich meine Frage vorher beantwortet haben.

Mara''',None,()),
    ('19_Buero_Antwort_2025.eml','2025-05-22T09:40:00+02:00','Hannes Riedel <buero@mauerpark-musik.example>','Mara Noémi Zwirn <mara.zwirn@klavierraum.example>','Re: Formular noch nicht freigegeben', '''Hallo Mara,

danke, ich spreche Kunibert darauf an. Der Portalvermerk kommt aus dem Import und ist keine Nachricht an eine Behörde. Ich dachte beim Kopieren, „final“ sei die abgegebene Fassung. Ich kann die Rentenfrage nicht beantworten. Die Information, die wir bekommen haben, sagt nur, dass die Abrechnung vorerst weiter über Honorar laufen kann.

Ich habe die Datei nicht gelöscht, damit wir den Stand wiederfinden. Die Buchhaltung bekommt zunächst die Liste; ich werde daneben „Rückfrage offen“ vermerken. Bitte schick die Konzertprogramme für Juni trotzdem bald. Der Druckauftrag geht am Dienstag raus und ich möchte deine Schüler nicht aus dem Programm nehmen.

Viele Grüße
Hannes''','18_Zustimmung_Nachfrage_2025.eml',()),
    ('20_Vertretung_abgelehnt.eml','2025-11-04T12:08:00+01:00','Hannes Riedel <buero@mauerpark-musik.example>','Mara Noémi Zwirn <mara.zwirn@klavierraum.example>','Lea am Donnerstag', '''Hallo Mara,

für Donnerstag kann ich Lea derzeit nicht ins Haus lassen. Das bei uns gespeicherte Führungszeugnis ist älter als die nach unserer Hausregel zulässigen drei Jahre. Bitte schick entweder den neueren Nachweis oder biete den Familien Ersatztermine an. Ihre fachliche Qualifikation steht nicht infrage; die Kolleginnen waren mit der Vertretung im Oktober 2023 zufrieden.

Ich habe noch keine Absagen an die Eltern verschickt. Wenn du bis 16 Uhr keine andere Lösung hast, informiere ich sie. Kunibert möchte nicht, dass jemand kurzfristig einspringt, den wir nie gesehen haben. Das gilt auch dann, wenn du die Vertretung selbst bezahlst. Für deinen ausgefallenen Unterricht würde sonst kein Honorar anfallen, wie im Vertrag vereinbart.

Hannes''',None,()),
    ('21_Korrektur_Stundenliste.eml','2025-02-07T15:18:00+01:00','Nils Brandt <buchhaltung@mauerpark-musik.example>','Mara Noémi Zwirn <mara.zwirn@klavierraum.example>','Bitte neue Januarrechnung', '''Hallo Frau Zwirn,

bitte senden Sie eine Ersatzrechnung über 2.448 Euro. Nach Hannes' Liste sind im Januar 64 Unterrichtseinheiten und vier vergütete Absagen zu berücksichtigen. Ihre bisherige Rechnung über 2.592 Euro ist noch nicht bezahlt. Ich habe den vorbereiteten Überweisungsauftrag vor Freigabe entfernt.

Die vier streitigen Absagen lassen sich in meiner Buchhaltung nicht aufklären. Bitte besprechen Sie die Ersatztermine direkt mit Hannes. Wenn sich daraus eine weitere Zahlung ergibt, benötige ich einen gesonderten Beleg. Schreiben Sie bitte auf die neue Rechnung, dass sie die alte ersetzt; sonst stehen beim nächsten Export zwei Januarumsätze nebeneinander.

Die Notenauslage über 27,80 Euro läuft unabhängig davon über das Auslagenkonto. Das ist kein Unterrichtshonorar und soll nicht nochmals in der Rechnung stehen.

Freundliche Grüße
Nils Brandt''',None,()),
    ('22_Geschaeftsfuehrung_Termin.eml','2026-09-28T09:12:00+02:00','Kunibert Schnurr <ks@mauerpark-musik.example>','Mara Noémi Zwirn <mara.zwirn@klavierraum.example>','Gespräch am 6. Oktober', '''Liebe Frau Zwirn,

wir halten den Termin am 6. Oktober um 10 Uhr frei. Sie können gern eine beratende Person mitbringen. Ich habe Hannes gebeten, Ihnen alle zu Ihrer Person gespeicherten Fassungen zu kopieren. Bislang ist keine Nachforderung gegen uns ergangen. Das Schreiben vom September fordert Unterlagen an; wir müssen die Fragen dennoch sorgfältig beantworten.

Bitte geben Sie bis dahin keine neue Erklärung nur deshalb ab, weil im Portal ein Feld offen ist. Wir werden Ihre Frage zur früheren Datei mitprüfen lassen. Ich kann Ihnen heute weder eine bestimmte Vertragsform zusagen noch Ihre persönliche Beitragsfrage beantworten. Über eine mögliche Anstellung können wir sprechen, sobald Umfang und Kosten geklärt sind.

Die Oktobertermine bleiben für Ihre Schüler reserviert. Eine Entscheidung über die künftige Zusammenarbeit wollen wir nicht während des laufenden Unterrichtsflurs treffen.

Mit freundlichen Grüßen
Kunibert Schnurr
Geschäftsführer''',None,())]
    for row in mails:mail(p,*row)
    raw(p,'23_Chat_Klavierteam_Juni2025.txt','''WhatsApp-Export „Klavierteam“ · Auszug 04.06.–10.06.2025 · exportiert von Mara Noémi Zwirn
04.06.2025, 20:52 - Mara: Hat jemand am 5. Juli Raum 2? Ich muss zwei Donnerstagsstunden wegen einer Konzertreise verlegen.
04.06.2025, 21:01 - Lea: Ich habe dort nichts. Aber im Kalender steht Sommerfest.
04.06.2025, 21:07 - Mara: Stimmt. Hannes sagte letzte Woche noch, Samstag ginge.
05.06.2025, 09:11 - Hannes: Bitte erst mit mir klären. Ich sehe die Elternzusagen, ihr nur die Raumbuchungen.
05.06.2025, 11:33 - Mara: Die Reise bleibt. Ich kann nach Rückkehr am Donnerstagabend zwei Doppelstunden machen. Das ist keine Frage nach Urlaub, ich suche einen Ersatzraum.
05.06.2025, 12:06 - Hannes: Ich weiß. Der Knopf heißt leider so. Für die Mappenschüler muss das trotzdem mit dem Probespiel zusammenpassen.
06.06.2025, 17:24 - Mara: Familie H. kann Sonntag, die anderen noch offen.
09.06.2025, 15:48 - Hannes: Jetzt freigegeben. Ich habe Raum 4 für Sonntag eingetragen. Elterninformation geht über dich, bitte Kopie ans Büro.
10.06.2025, 08:19 - Lea: Ich könnte Sonntag eine Stunde übernehmen.
10.06.2025, 08:42 - Mara: Danke, diesmal mache ich selbst. Das Geld für Vertretung frisst mir sonst fast das ganze Honorar weg.
Exporthinweis des Geräts: Medien wurden beim Export nicht aufgenommen.''')
    raw(p,'24_Telefonnotiz_20260923.txt','''23.09.2026, 13:05 Uhr – Anruf Hannes bei Mara, von Mara nach dem Gespräch notiert
Hannes: Prüferin fragt nach Verträgen und Stundenplänen. Er könne mir das Schreiben schicken. Es betreffe „eigentlich nur den Betrieb“. Auf meine Frage nach persönlichen Beiträgen sagte er, das müsse Kunibert mit dem Steuerbüro klären.
Ich: Ich hatte doch damals gefragt, ob der Sparplan reicht. Seitdem habe ich nichts mehr gehört.
Hannes: Er habe die Rückfrage mit auf die Liste gesetzt. Er dachte, ich hätte später mit Kunibert gesprochen. Er wisse nicht, wer den Vermerk „erhalten“ verändert hat; wahrscheinlich der Import.
Ich: Bitte alles schicken, auch das Formular. Nichts neu unterschreiben lassen, bevor ich es geprüft habe.
Hannes: Kein Problem. Er ruft nächste Woche wegen des Termins an. Die vier Schüler für die Vorbereitung soll ich weiter einplanen.
Ende etwa 13:17. Keine Tonaufnahme. Die Sätze sind aus meiner Erinnerung, nicht wörtlich mitgeschrieben.''')
    raw(p,'25_Chat_Eigene_Schueler.txt','''Signal-Export „Samstagsunterricht“ · Mara Noémi Zwirn und Paul Werner
08.01.2025, 18:05 - Paul: Bleibt es bei 48 Euro für die Stunde bei dir?
08.01.2025, 18:22 - Mara: Ja, für 60 Minuten. Ich stelle Ende des Monats die Rechnung direkt an dich.
08.01.2025, 18:31 - Paul: Samstag um zehn wäre gut. Was passiert, wenn ich wegen Schichtdienst nicht kann?
08.01.2025, 18:40 - Mara: Bis Freitag zehn Uhr kannst du absagen. Danach berechne ich die Stunde, außer wir finden in derselben Woche Ersatz. Bei meinen eigenen Absagen natürlich kein Honorar.
16.05.2025, 09:14 - Paul: Du machst doch auch Unterricht in der Akademie. Kann ich dort einfach in deinen Raum kommen?
16.05.2025, 09:52 - Mara: Nein. Unser Vertrag ist privat bei mir zu Hause. Die Akademieräume kann ich nicht einfach mitbenutzen.
20.09.2026, 13:18 - Paul: Rechnung ist bezahlt, danke. Im Oktober bin ich zwei Wochen weg.
20.09.2026, 14:03 - Mara: Dann planen wir nur den 3. und 24. Oktober. Ich reserviere nichts für die beiden anderen Samstage.
Geräteexport 24.09.2026. Keine Anlagen.''')
    csvfile(p,'26_Honorarkonto_2025.csv',[['Monat','Unterricht_Einheiten','bezahlte_Ausfälle','Veranstaltung_Einheiten','Satz_EUR','Honorar_EUR','Buchungsvermerk']]+[[f'2025-{i:02}',u,a,v,36,(u+a+v)*36,n] for i,(u,a,v,n) in enumerate([(64,4,0,'Januar Korrektur MS2025-01K'),(60,3,2,'Probespiel72EUR separat'),(72,4,0,''),(48,2,0,'Osterferien'),(68,4,0,''),(56,2,2,'Konzertreise und Nachholstunden'),(36,1,0,'Sommerferien ab24.07.'),(0,0,0,'Sommerferien'),(76,3,0,''),(64,2,0,'Herbstferien'),(72,2,0,'Keine Zahlung für eigene Ausfälle'),(48,3,2,'Schülerkonzert')],1)])
    return p


if __name__=='__main__':
    music();programmer()
    (QA/'native-inventar.json').write_text(json.dumps(inventories,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:len(v) for k,v in inventories.items()},ensure_ascii=False))
