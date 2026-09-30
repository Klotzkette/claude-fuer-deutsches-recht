#!/usr/bin/env python3
"""Gezielte, nur lesende Regressionen für Plugin und Mainzer Originalakte.

Mit --qa-dir werden zusätzlich die tatsächlich gerenderten PDF/PNG geprüft.
Die Sichtprüfung ist separat dokumentiert; diese Tests bewerten keine Modellantwort.
"""

from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import date
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
import hashlib
import json
from pathlib import Path
import re
import sys
import unittest
from urllib.parse import unquote, urlsplit

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageChops
from pypdf import PdfReader
import yaml

from quality_lab import validate_profile

ROOT = Path(__file__).resolve().parents[1]
SLUG = "ki-verordnung-transparenzpruefer"
PLUGIN = ROOT / SLUG
CASE = ROOT / "testakten/ki-transparenz-kanzlei-kommunikation-mainz"
FIXTURE = ROOT / "scripts/fixtures/ki-transparenz-mainz"
PROFILE = json.loads((ROOT / "quality/evals" / f"{SLUG}.json").read_text())
QA = json.loads((ROOT / "quality" / f"{SLUG}-native-qa.json").read_text())
SPECS = json.loads((FIXTURE / "documents.json").read_text())["documents"]
ORIGINALS = sorted(CASE.glob("[0-9][0-9]_*"))
SKILLS = sorted((PLUGIN / "skills").glob("*/SKILL.md"))
QA_DIR = None
DISCLAIMER = "Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr."
DISCLAIMER_EN = "This test case file was generated with AI and is an experiment. Use at your own responsibility and risk."


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def compact(text):
    return re.sub(r"\s+", "", text).replace("\u00ad", "")


def content(path):
    if path.suffix == ".docx":
        return "\n".join(p.text for p in Document(path).paragraphs)
    if path.suffix == ".pdf":
        return "\n".join(p.extract_text() for p in PdfReader(path).pages)
    if path.suffix == ".eml":
        return BytesParser(policy=policy.default).parsebytes(path.read_bytes()).get_content()
    return path.read_text(encoding="utf-8")


class TransparenzRegression(unittest.TestCase):
    def test_manifest(self):
        manifest = json.loads((PLUGIN / ".claude-plugin/plugin.json").read_text())
        self.assertEqual(manifest["name"], SLUG)
        self.assertRegex(manifest["version"], r"^\d+\.\d+\.\d+$")
        self.assertGreaterEqual(tuple(map(int, manifest["version"].split("."))), (445, 11, 0))
        self.assertLessEqual(len(manifest["description"]), 300)
        self.assertEqual(manifest["license"], "Apache-2.0 OR MIT")
        self.assertEqual(manifest["author"], {"name": "Klotzkette", "email": "39582916+Klotzkette@users.noreply.github.com"})

    def test_skill_contracts(self):
        self.assertEqual(len(SKILLS), 8)
        for path in SKILLS:
            with self.subTest(skill=path.parent.name):
                raw = path.read_text()
                header, body = raw.split("---", 2)[1:]
                meta = yaml.safe_load(header)
                self.assertEqual(set(meta), {"name", "description"})
                self.assertEqual(meta["name"], path.parent.name)
                self.assertRegex(meta["name"], r"^[a-z0-9-]{1,64}$")
                self.assertTrue(80 <= len(meta["description"]) <= 1024)
                self.assertEqual(re.findall(r"^## (\d+)\.", body, re.M), list("123456"))
                output = body.split("## 5.")[1].split("## 6.")[0]
                for requirement in ("Ausformulierungspflicht", "vollständigen", "Skelette", "Times New Roman 11 pt", "dezimal"):
                    self.assertIn(requirement, output)
                self.assertGreater(len(body.split()), 300)

    def test_portable_skill_sources(self):
        for path in SKILLS:
            links = re.findall(r"\]\(([^)]+)\)", path.read_text())
            self.assertGreaterEqual(len(links), 2)
            for link in links:
                parsed = urlsplit(link)
                if parsed.scheme:
                    self.assertEqual(parsed.scheme, "https")
                    self.assertTrue(parsed.hostname.endswith(".europa.eu") or parsed.hostname == "www.gesetze-im-internet.de")
                else:
                    resolved = (path.parent / unquote(parsed.path)).resolve()
                    self.assertTrue(resolved.is_relative_to(PLUGIN.resolve()), link)
                    self.assertTrue(resolved.is_file(), link)

    def test_editorial_style(self):
        for path in PLUGIN.rglob("*.md"):
            raw = path.read_text()
            with self.subTest(path=str(path.relative_to(PLUGIN))):
                self.assertNotIn("**", raw)
                self.assertNotIn(chr(167), raw)
                self.assertNotRegex(raw, re.compile(r"^#{1,6}\s+(?:[IVX]+|[A-Za-z])[.)]\s", re.M))
        mini = PLUGIN / f"{SLUG}-schnellstart.md"
        self.assertLessEqual(len(mini.read_bytes()), 7500)
        self.assertGreater(len(mini.read_text()), 5000)

    def test_local_eval_schema_and_hashes(self):
        validate_profile(PROFILE, SLUG, PLUGIN, ROOT)
        self.assertEqual(PROFILE["prompt_editorial_review"]["method"], "desk_review")
        self.assertEqual(PROFILE["prompt_workflow_review"]["method"], "desk_review")
        for kind, key in (("werkstatt", "workshop_review"), ("schnellstart", "mini_review")):
            self.assertEqual(PROFILE[key]["sha256"], digest(PLUGIN / f"{SLUG}-{kind}.md"))

    def test_workflow_coverage(self):
        cases = {c["id"]: c for c in PROFILE["cases"]}
        self.assertEqual(len(cases), 15)
        self.assertEqual({c["target_skill"] for c in cases.values()}, {p.parent.name for p in SKILLS})
        complete = cases["ordnerauftrag-bis-enddokument"]
        self.assertEqual({Path(p).name for p in complete["input_files"]}, {p.name for p in ORIGINALS})
        self.assertEqual(cases["leerer-start-kein-faktendump"]["input_files"], [])
        self.assertIn("Fortsetzung", cases["redaktion-nachgereichte-antwort"]["request"])
        self.assertGreaterEqual(len(PROFILE["prompt_workflow_review"]["opening_modes"]), 3)

    def test_verified_source_and_legal_regressions(self):
        source = (PLUGIN / "references/rechtsstand-artikel-50.md").read_text()
        for token in ("2026/1744", "27. Juli 2026", "2. Dezember 2026", "Artikel 111 Absatz 4", "COM(2025) 837", "laufendes Verfahren", "Artikel 4a", "keine Deepfake", "unverbindlich", "freiwillig", "nichtgerichtlicher"):
            self.assertIn(token, source + json.dumps(PROFILE, ensure_ascii=False))
        for section in range(1, 6):
            self.assertIn(f"### 3.{section}.", source)
        for token in ("angestellter Rechtsanwalt", "nichtöffentlich", "menschlicher Überprüfung oder redaktioneller Kontrolle", "gesonderte Anbieterprüfung nach Absatz 2", "erste", "biometrischer"):
            self.assertIn(token, source)
        for record in PROFILE["sources"]:
            host = urlsplit(record["url"]).hostname
            self.assertTrue(host.endswith(".europa.eu") or host == "www.gesetze-im-internet.de")
            self.assertIn(record["checked_on"], {"2026-09-28", "2026-09-30"})

    def test_final_guideline_edge_case_coverage(self):
        required = {
            "email-agent-oder-menschliche-antwort": "chat-und-telefonhinweise",
            "internes-gutachten-keine-b2b-pauschalausnahme": "anbieterkennzeichnung-pruefen",
            "erfundener-sprecher-interne-vorfuehrung": "deepfake-offenlegung",
            "markierung-ohne-detektionsweg": "anbieterkennzeichnung-pruefen",
            "altinhalt-nicht-systemuebergang": "rollen-und-anwendungsdatum",
            "erneute-textpruefung-und-clip-ohne-vorspann": "transparenzfreigabe-dokumentieren",
        }
        cases = {case["id"]: case for case in PROFILE["cases"]}
        for case_id, skill in required.items():
            with self.subTest(case=case_id):
                case = cases[case_id]
                self.assertEqual(case["target_skill"], skill)
                self.assertGreaterEqual(len(case["criteria"]), 3)
                self.assertEqual(case["input_files"], [])
                for criterion in case["criteria"]:
                    self.assertEqual(criterion["deliverables"], case["deliverables"])
        self.assertIn("keine behaupteten bestandenen Modellläufe", PROFILE["prompt_workflow_review"]["limits"])

    def test_both_plugins_ship_final_guidelines(self):
        references = {
            SLUG: "rechtsstand-artikel-50.md",
            "ki-vo-ai-act-pruefer": "artikel-50-leitlinien-2026.md",
        }
        for slug, filename in references.items():
            with self.subTest(plugin=slug):
                plugin = ROOT / slug
                source = (plugin / "references" / filename).read_text()
                self.assertIn("https://ec.europa.eu/newsroom/dae/redirection/document/131215", source)
                self.assertIn("153", source)
                self.assertIn("154", source)
                self.assertIn("Artikel 111 Absatz 4", source)
                profile = json.loads((ROOT / "quality/evals" / f"{slug}.json").read_text())
                validate_profile(profile, slug, plugin, ROOT)
                for kind, key in (("werkstatt", "workshop_review"), ("schnellstart", "mini_review")):
                    path = plugin / f"{slug}-{kind}.md"
                    text = path.read_text()
                    self.assertIn("131215", text)
                    self.assertIn("154", text)
                    self.assertEqual(profile[key]["sha256"], digest(path))
                    if kind == "schnellstart":
                        self.assertLessEqual(len(path.read_bytes()), 7500)
                official = [s for s in profile["sources"] if s["url"].endswith("/131215")]
                self.assertEqual(len(official), 1)
                self.assertEqual(official[0]["checked_on"], "2026-09-30")

    def test_short_form_keeps_statutory_quality_standard(self):
        mini = (PLUGIN / f"{SLUG}-schnellstart.md").read_text()
        for criterion in ("Wirksamkeit", "Interoperabilität", "Robustheit", "Zuverlässigkeit", "Machbarkeit", "Inhaltsspezifika", "Kosten", "Stand der Technik"):
            self.assertIn(criterion, mini)
        reference = (ROOT / "ki-vo-ai-act-pruefer/references/artikel-50-leitlinien-2026.md").read_text()
        self.assertIn("unabhängig von direkter Interaktion", reference)

    def test_original_inventory_and_rubric(self):
        self.assertEqual(len(ORIGINALS), 12)
        self.assertEqual(Counter(p.suffix for p in ORIGINALS), {".docx": 3, ".pdf": 3, ".eml": 4, ".csv": 1, ".txt": 1})
        self.assertEqual([p.name[:2] for p in ORIGINALS], [f"{i:02}" for i in range(1, 13)])
        rubric = yaml.safe_load((CASE / "rubric.yaml").read_text())
        self.assertEqual(rubric["plugin"], SLUG)
        for check in rubric["checks"]:
            self.assertIn(check["check_type"], {"file_exists", "human_review"})
            if check["check_type"] == "file_exists":
                self.assertTrue((CASE / check["path"]).is_file())

    def test_email_integrity(self):
        bodies, ids = set(), set()
        for path in CASE.glob("*.eml"):
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            self.assertFalse(message.defects)
            for key in ("From", "To", "Date", "Subject", "Message-ID"):
                self.assertTrue(message[key], (path.name, key))
            self.assertIsNotNone(parsedate_to_datetime(message["Date"]).utcoffset())
            self.assertRegex(str(message["Date"]), r"\d{2}:\d{2}:\d{2} [+-]\d{4}$")
            for _, address in getaddresses([str(message[k]) for k in ("From", "To", "Cc") if message[k]]):
                self.assertTrue(address.endswith(".example"), address)
            body = message.get_content()
            self.assertGreater(len(body.split()), 200)
            bodies.add(body)
            ids.add(str(message["Message-ID"]))
        self.assertEqual(len(bodies), 4)
        self.assertEqual(len(ids), 4)

    def test_csv_integrity(self):
        with next(CASE.glob("*.csv")).open(encoding="utf-8", newline="") as stream:
            reader = csv.DictReader(stream, delimiter=";")
            rows = list(reader)
            self.assertEqual(len(reader.fieldnames), 11)
        self.assertEqual(len(rows), 9)
        self.assertEqual(len({r["kennung"] for r in rows}), 9)
        for row in rows:
            date.fromisoformat(row["einsatzbeginn"])
            self.assertNotIn(None, row)
            for value in row.values():
                self.assertTrue(value)
                self.assertFalse(value.startswith(("=", "+", "-", "@")))

    def test_native_content_matches_fixtures(self):
        self.assertEqual(len(SPECS), 6)
        for spec in SPECS:
            actual = compact(content(CASE / spec["file"]))
            for text in [spec["title"], spec["author"], spec["recipient"], spec["date"], spec["reference"]]:
                self.assertIn(compact(text), actual, spec["file"])
            for section in spec["sections"]:
                self.assertRegex(section["heading"], r"^\d+\.")
                self.assertIn(compact(section["heading"]), actual)
                for paragraph in section["paragraphs"]:
                    self.assertIn(compact(paragraph), actual, spec["file"])

    def test_docx_style(self):
        for path in CASE.glob("*.docx"):
            doc = Document(path)
            for name in ("Normal", "Title", "Heading 1"):
                style = doc.styles[name]
                self.assertEqual(style.font.name, "Times New Roman")
                self.assertEqual(style.font.size.pt, 11)
                self.assertEqual(str(style.font.color.rgb), "000000")
                self.assertFalse(list(style.element.iter(qn("w:pBdr"))))
            self.assertEqual(doc.paragraphs[0].style.name, "Title")
            self.assertIn("PAGE", doc.sections[0].footer._element.xml)
            self.assertFalse(doc.inline_shapes)

    def test_pdf_and_image_payload(self):
        for path in CASE.glob("*.pdf"):
            reader = PdfReader(path, strict=True)
            self.assertEqual(len(reader.pages), 2 if path.name.startswith("10_") else 1)
            for number, page in enumerate(reader.pages, 1):
                self.assertIn(f"Seite {number}", page.extract_text())
                sizes = []
                page.extract_text(visitor_text=lambda text, cm, tm, font, size: sizes.append(size) if text.strip() else None)
                self.assertGreaterEqual(min(sizes), 11)
        picture = PdfReader(next(CASE.glob("10_*.pdf"))).pages[1]
        self.assertEqual(len(picture.images), 1)
        self.assertGreater(picture.images[0].image.width, 500)

    def test_disclaimers_and_no_solution_in_originals(self):
        for path in ORIGINALS:
            text = content(path)
            self.assertNotIn(DISCLAIMER, text)
            self.assertNotIn(DISCLAIMER_EN, text)
            self.assertNotIn(chr(167), text)
            self.assertNotRegex(text.lower(), r"musterlösung|bewertungsschlüssel|instructor|placeholder")
        for name in ("README.md", "README.txt"):
            raw = (CASE / name).read_text()
            self.assertIn(DISCLAIMER, raw)
            self.assertIn(DISCLAIMER_EN, raw)
        self.assertTrue((CASE / "README.txt").read_text().startswith(DISCLAIMER))

    def test_visual_record_matches_originals(self):
        self.assertEqual(QA["document_count"], 12)
        self.assertEqual(QA["page_count"], 17)
        self.assertEqual({d["file"] for d in QA["documents"]}, {p.name for p in ORIGINALS})
        self.assertEqual(sum(len(d["pages"]) for d in QA["documents"]), 17)
        for record in QA["documents"]:
            self.assertEqual(record["sha256"], digest(CASE / record["file"]), record["file"])
            self.assertEqual([p["page"] for p in record["pages"]], list(range(1, len(record["pages"]) + 1)))

    def test_rendered_native_formats(self):
        if QA_DIR is None:
            self.skipTest("Renderprüfung benötigt --qa-dir mit den geprüften Seitenbildern")
        for record in QA["documents"]:
            original = CASE / record["file"]
            folder = QA_DIR / original.stem
            self.assertEqual(len(list(folder.glob("page-*.png"))), len(record["pages"]))
            pdf = original if original.suffix == ".pdf" else next(folder.glob("*.pdf"))
            reader = PdfReader(pdf)
            self.assertEqual(len(reader.pages), len(record["pages"]))
            for page in record["pages"]:
                path = folder / f"page-{page['page']}.png"
                self.assertEqual(digest(path), page["png_sha256"], str(path))
                with Image.open(path) as bitmap:
                    self.assertGreaterEqual(bitmap.width, 1000)
                    difference = ImageChops.difference(bitmap.convert("RGB"), Image.new("RGB", bitmap.size, "white"))
                    bounds = difference.getbbox()
                    self.assertIsNotNone(bounds)
                    self.assertGreater(bounds[0], 25)
                    self.assertGreater(bounds[1], 25)
                    self.assertLess(bounds[2], bitmap.width - 25)
                    self.assertLess(bounds[3], bitmap.height - 25)
            extracted = compact("\n".join(re.sub(r"(?m)^Seite \d+\s*$", "", p.extract_text()) for p in reader.pages))
            if original.suffix in {".docx", ".eml", ".txt"}:
                self.assertIn(compact(content(original)), extracted, record["file"])
            if original.suffix == ".csv":
                with original.open(encoding="utf-8", newline="") as stream:
                    for row in csv.DictReader(stream, delimiter=";"):
                        for value in row.values():
                            self.assertIn(compact(value), extracted)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--qa-dir", type=Path)
    args, remaining = parser.parse_known_args()
    QA_DIR = args.qa_dir
    unittest.main(argv=[sys.argv[0], *remaining], verbosity=2)
