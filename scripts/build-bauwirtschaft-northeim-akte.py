#!/usr/bin/env python3
"""Vergabeakte Northeim. Autor: Klotzkette. Nur fallbezogene Ausgaben."""
from pathlib import Path
from datetime import datetime
from decimal import Decimal
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
import csv
import hashlib
import json
import os
import sys
import textwrap

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from PIL import Image, ImageDraw, PngImagePlugin
from akten_build_runtime import screen_font, serif_font_path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
SLUG = 'bauwirtschaft-vergabeverfahren-feuerwehrhaus-northeim'
CASE = ROOT / 'testakten' / SLUG
QA = Path('/tmp/bauwirtschaft-assets/.northeim-qa')
AKTE = 'NO-FH26-L430'
CITY = 'Stadt Northeim | Zentrale Vergabestelle'
ALPHA = 'Leinetal Lufttechnik GmbH'
BETA = 'Weserklima Anlagenbau GmbH'
GAMMA = 'Harzraum Gebäudetechnik GmbH'
STAMP = datetime(2026, 9, 25, 14, 0)

# Mengengerüst, ausgeschriebene Leistungsmerkmale und drei eigenständige Kalkulationen.
POSITIONS = [
 ['01.010','Baustelleneinrichtung',1,'psch','Einrichten und Räumen; Lagerfläche 25 m²; Hebezeuge, Schutz und Entsorgung der eigenen Verpackungen.',12800,14200,14900],
 ['01.020','Zentrales Lüftungsgerät',1,'St','Zu-/Abluft je 6.000 m³/h bei 600 Pa extern; Wärmerückgewinnung mindestens 80 % trocken; SFP höchstens 1,80 kW/(m³/s); Gehäuse höchstens 3.200 x 1.600 x 1.800 mm; BACnet/IP.',78000,102000,107500],
 ['01.030','Dachabluftventilator',2,'St','Je 2.200 m³/h bei 350 Pa; EC-Antrieb, Dachaufsatz, Revisionsschalter, elastischer Anschluss; betriebsfertig.',8600,9400,9800],
 ['01.040','Rechteckkanal verzinkt',560,'m²','Kanalblech einschließlich Formstücke und Verbindungen; längste Kante bis 1.000 mm; dicht gefügt; Abrechnung nach Kanaloberfläche.',118,126,131],
 ['01.050','Rundrohrnetz',210,'m','Wickelfalzrohr DN 160 bis DN 315, einschließlich Bögen, Abzweige und Dichtungen; gemessene Rohrmittellänge.',76,81,84],
 ['01.060','Wärmedämmung Kanalnetz',460,'m²','Mineralfaserdämmung 40 mm mit geschlossener Außenkaschierung; Stöße verklebt, Befestigung ohne offene Fugen.',41,44,46],
 ['01.070','Brandschutzklappen',14,'St','Motorisch, 24 V, zwei Endlagenkontakte; zugelassen für jeweilige Wandart, lichte Größen 200 x 200 bis 600 x 400 mm; Einbau mit Nachweis.',870,930,965],
 ['01.080','Volumenstromregler',18,'St','DN 160 bis DN 250, einstellbar 100 bis 900 m³/h, Stellantrieb 24 V, zugängliche Messanschlüsse.',395,430,450],
 ['01.090','Luftdurchlässe',42,'St','Decken- oder Wanddurchlass mit Anschlusskasten und Drossel, pulverbeschichtet RAL 9010; 80 bis 450 m³/h je Durchlass.',155,168,172],
 ['01.100','Kulissenschalldämpfer',6,'St','Einfügungsdämpfung mindestens 20 dB bei 250 Hz; einschließlich Anschlussübergängen; Druckverlust höchstens 50 Pa.',780,825,860],
 ['01.110','Außenluftgitter',4,'St','Witterungsschutz und Insektengitter; freie Fläche je mindestens 0,45 m²; Rahmen und dauerhaft dichte Befestigung.',950,1020,1045],
 ['01.120','Regelung und Schaltschrank',1,'psch','Schaltschrank, Feldgeräte, CO₂-Regelung in vier Zonen, Frostschutz, BACnet/IP mit 40 Datenpunkten; Parametrierung und Schnittstellentest.',24500,26800,28100],
 ['01.130','Steuerleitungen',380,'m','Liefern und verlegen im eigenen Gewerk; halogenfrei, einschließlich Kleinmaterial und beidseitiger Beschriftung.',18,19.5,20],
 ['01.140','Befestigungen und Tragsystem',1,'psch','Sämtliche Kanal- und Gerätehalterungen einschließlich Schwingungsentkopplung und statischem Befestigungsnachweis.',12400,13700,14100],
 ['01.150','Brandschottungen',26,'St','Durchführungen bis 0,12 m², zur Wand passende Abschottung einschließlich Kennzeichnung und Fotodokumentation.',185,198,205],
 ['01.160','Inbetriebnahme',1,'psch','Funktionsprüfung, Einregulierung der Anlage und Sicherheitsketten, Anwesenheit bei integrierter Funktionsprüfung an zwei Arbeitstagen.',9200,9800,10200],
 ['01.170','Volumenstrommessung',42,'St','Messung an jedem Durchlass, Soll-/Ist-Protokoll, Nachregulierung auf höchstens 10 % Abweichung.',85,92,96],
 ['01.180','Bestandsdokumentation',1,'psch','Revisionspläne PDF und bearbeitbar, Betriebshandbuch, Geräteliste und Messprotokolle; zwei Papierausfertigungen.',4600,4950,5100],
 ['01.190','Einweisung Betriebspersonal',2,'Termin','Je vier Stunden für bis zu acht Personen einschließlich Vorbereitung und Teilnehmernachweis.',650,700,740],
 ['01.200','Erste Wartung',1,'psch','Ein Wartungstermin zwölf Monate nach Abnahme, Filterwechsel und Funktionsprüfung; Verbrauchsfilter enthalten.',3200,3450,3600],
]
BIDDERS = [
 dict(id='08',name=ALPHA,short='leinetal',date='28.08.2026',person='Lena Winter',address='Am Werkbogen 8, 37154 Northeim',device='Auenluft AL 6000 K',price_col=5),
 dict(id='10',name=BETA,short='weserklima',date='28.08.2026',person='Martin Heuer',address='Gewerbering 26, 37603 Holzminden',device='Flusswind FW 6200',price_col=6),
 dict(id='12',name=GAMMA,short='harzraum',date='31.08.2026',person='Sven Martens',address='Am Technikfeld 14, 37520 Osterode am Harz',device='Nordvent NV 6000',price_col=7),
]

def money(value):
    return f'{value:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

def native(text):
    return text.replace('.example', '.de').replace('fallinternen Geschäftszeichen', 'Geschäftszeichen')

def total(b):
    return sum(Decimal(str(p[2])) * Decimal(str(p[b['price_col']])) for p in POSITIONS)

def paragraphs(text):
    return [p.strip() for p in textwrap.dedent(text).strip().split('\n\n') if p.strip()]

def docx(name, org, title, date, sections, closing):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.1)
    for key in ['Normal','Title','Heading 1','Heading 2']:
        style = doc.styles[key]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(11 if key == 'Normal' else 15 if key == 'Title' else 12)
        style.font.color.rgb = RGBColor(0,0,0)
        style.paragraph_format.space_after = Pt(7)
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if 'Theme' in attr or 'theme' in attr:del fonts.attrib[attr]
        for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),'Times New Roman')
    for border in doc.styles.element.iter(qn('w:pBdr')):
        border.getparent().remove(border)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.08
    doc.core_properties.author = 'Klotzkette'
    doc.core_properties.last_modified_by = 'Klotzkette'
    doc.core_properties.created = doc.core_properties.modified = STAMP
    doc.core_properties.title = title
    doc.add_paragraph(native(org))
    doc.add_paragraph(f'{date} | {AKTE}')
    doc.add_paragraph(title, 'Title')
    for heading, body in sections:
        if heading == 'PAGE':
            doc.add_page_break()
            continue
        if heading:
            doc.add_paragraph(heading, 'Heading 1')
        for p in paragraphs(body):
            doc.add_paragraph(native(p))
    doc.add_paragraph(native(closing))
    footer = sec.footer.paragraphs[0]
    footer.add_run(f'{AKTE} | Seite ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.save(CASE / name)

def pdf(name, org, title, date, pages):
    if 'TNR' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('TNR',str(serif_font_path())))
        pdfmetrics.registerFont(TTFont('TNR-Bold',str(serif_font_path(bold=True))))
    body = ParagraphStyle('body', fontName='TNR', fontSize=11, leading=14, spaceAfter=9)
    head = ParagraphStyle('head', parent=body, fontName='TNR-Bold', fontSize=12, spaceBefore=6)
    title_style = ParagraphStyle('title', parent=head, fontSize=16, leading=19, spaceAfter=15)
    flows = []
    for i, sections in enumerate(pages):
        if i: flows.append(PageBreak())
        flows += [Paragraph(escape(org), body), Paragraph(escape(f'{date} | {AKTE}'), body)]
        if i == 0: flows.append(Paragraph(escape(title), title_style))
        for h, text in sections:
            if h: flows.append(Paragraph(escape(h), head))
            flows += [Paragraph(escape(native(p)).replace('\n','<br/>'), body) for p in paragraphs(text)]
    def footer(c, d):
        c.setFont('TNR', 9)
        c.drawString(60, 30, f'{AKTE} | {title}')
        c.drawRightString(535, 30, str(d.page))
    SimpleDocTemplate(str(CASE/name), pagesize=A4, leftMargin=60, rightMargin=60, topMargin=48, bottomMargin=52, title=title, author='Klotzkette').build(flows, onFirstPage=footer, onLaterPages=footer)

def eml(name, sender, to, iso, subject, body, token, reply=None):
    msg = EmailMessage(policy=SMTP)
    msg['From'], msg['To'], msg['Date'] = native(sender), native(to), format_datetime(datetime.fromisoformat(iso))
    msg['Subject'] = f'{AKTE} | {subject}'
    msg['Message-ID'] = f'<{token}@portal.northeim-vergabe.de>'
    msg['X-Case-Reference'] = AKTE
    msg['Received'] = f'from portal.northeim-vergabe.de by mail.aktenpost.de with ESMTP id {token}; {msg["Date"]}'
    if reply:
        msg['In-Reply-To'] = f'<{reply}@portal.northeim-vergabe.de>'
        msg['References'] = f'<{reply}@portal.northeim-vergabe.de>'
    msg.set_content(native(textwrap.dedent(body)).strip() + '\n', charset='utf-8')
    msg.add_alternative('<html><body>' + ''.join('<p>'+escape(native(p)).replace('\n','<br>')+'</p>' for p in paragraphs(body)) + '</body></html>', subtype='html', charset='utf-8')
    (CASE/name).write_bytes(msg.as_bytes())

def screens():
    for filename, role, account, stamp, title, rows in [
        ('15_portal_bieter.png','Bieterzugang','Lena Winter · Leinetal Lufttechnik GmbH','31.08.2026 09:54:18 MESZ','Angebot erfolgreich eingereicht',[
            ('Vorgang',AKTE),('Empfangskennung','NO26-E001'),('Servereingang','31.08.2026 09:54:12 MESZ'),('Fristende','31.08.2026 10:00:00 MESZ'),('Paket','08_angebot_leinetal.xlsx + 09_eignung_leinetal.pdf'),('Dateien','2 Dateien · verschlüsselt gespeichert'),('Fassung','1 · keine Rücknahme'),('Status','Eingang bestätigt; Inhalt erst nach Fristablauf lesbar')]),
        ('16_portal_vergabestelle.png','Vergabestelle','Mara Seidel · Zentrale Vergabestelle','09.09.2026 12:06:31 MESZ','Kommunikation und Eingangskontrolle',[
            ('Vorgang',AKTE),('Nachricht','NO26-N018 · Nachweise und Aufklärung'),('Absendung','02.09.2026 09:00:00 MESZ'),('Empfänger','Leinetal Lufttechnik GmbH'),('Frist','09.09.2026 12:00:00 MESZ'),('Antwort','09.09.2026 11:41:02 MESZ · NO26-A019'),('Anhänge','19_antwort.docx / 20_referenz.pdf / 21_geraet.pdf'),('Arbeitsstatus','Eingang vollständig erfasst · fachliche Prüfung offen')])]:
        im = Image.new('RGB',(1600,1000),'#f4f6f7')
        d = ImageDraw.Draw(im)
        d.rectangle((0,0,1600,78),fill='#253d42')
        d.text((38,22),'Kommunales Vergabeportal',font=screen_font(28,True),fill='white')
        d.text((1100,27),role,font=screen_font(23),fill='white')
        d.rectangle((0,78,280,1000),fill='#e5eaeb')
        for y,s in [(120,'Übersicht'),(183,'Vergabeunterlagen'),(246,'Nachrichten'),(309,'Angebote'),(372,'Protokoll')]:
            d.text((30,y),s,font=screen_font(21),fill='#294047')
        d.text((322,110),account,font=screen_font(22),fill='#35434b')
        d.text((322,158),'Feuerwehrhaus Northeim · Los 430 Lüftung',font=screen_font(29,True),fill='#172e32')
        d.text((322,224),title,font=screen_font(26,True),fill='#285d42')
        for i,(k,v) in enumerate(rows):
            y=302+i*65
            d.line((322,y+46,1545,y+46),fill='#d4dcde',width=1)
            d.text((328,y),k,font=screen_font(20,True),fill='#37464b')
            d.text((572,y),v,font=screen_font(20),fill='#1e282b')
        d.text((322,917),stamp+' · Serverzeit Europe/Berlin',font=screen_font(20),fill='#58666b')
        info=PngImagePlugin.PngInfo();info.add_text('Author','Klotzkette')
        im.save(CASE/filename,pnginfo=info)

def main():
    CASE.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
    (QA/'build-data.json').write_text(json.dumps({'positions':POSITIONS,'bidders':BIDDERS,'totals':[float(total(b)) for b in BIDDERS]},ensure_ascii=False,indent=2),encoding='utf-8')
    a,b,c = [money(total(x)) for x in BIDDERS]
    docx('01_vergabe_einleitung.docx',CITY,'Einleitung der Bauvergabe','27.07.2026',[
        ('1 Vorhaben und Schätzgrundlage','Die Stadt errichtet das Feuerwehrhaus Northeim, Projektstandort Am Gerätehof 4, mit sechs Stellplätzen, Umkleiden, Werkstatt und Schulungsbereich. Der Standort wird in den Vergabeunterlagen projektbezogen bezeichnet. Der vorläufige Baubeginn ist Oktober 2026. Die Schätzung zum heutigen Versand der Bekanntmachung umfasst alle Bauleistungen einschließlich der ersten Wartung der technischen Anlagen. Grundstück und separat vergebene Planungsleistungen sind nicht enthalten.'),
        ('2 Gesamtwert und Lose','Der geschätzte Gesamtauftragswert beträgt 8.400.000,00 EUR netto. Er setzt sich aus Erd- und Gründungsarbeiten 820.000,00 EUR, Rohbau 2.120.000,00 EUR, Dach und Fassade 1.140.000,00 EUR, Ausbau 1.690.000,00 EUR, Heizung und Sanitär 760.000,00 EUR, Elektro 720.000,00 EUR, Lüftung 380.000,00 EUR sowie Außenanlagen 770.000,00 EUR zusammen. Die Kostenschätzung wurde durch den Fachbereich Hochbau am 24.07.2026 fortgeschrieben. Es sind keine weiteren Optionen vorgesehen.'),
        ('3 Verfahrensfestlegung','Der ab 01.01.2026 geltende Schwellenwert für Bauaufträge beträgt 5.404.000,00 EUR netto gemäß Artikel 1 der Delegierten Verordnung (EU) 2025/2152. Ausgangspunkt ist nach Paragraf 3 Absatz 7 VgV der Gesamtwert aller Lose. Für Los 430 wird die Ausnahme des Paragrafen 3 Absatz 9 VgV nicht genutzt. Das Lüftungslos wird daher im offenen EU-Verfahren nach VOB/A Abschnitt 2 vergeben. Eine isolierte Bewertung anhand des Loswerts von 380.000,00 EUR findet nicht statt.'),
        ('4 Organisation','Mara Seidel führt das Verfahren, Jonas Berg übernimmt die zweite Prüfung. Die fachliche Prüfung erfolgt durch Dipl.-Ing. Nora Brandt im Fachbereich Hochbau. Interessenkonflikte wurden von den drei Bearbeitenden verneint. Die Vergabestelle dokumentiert sämtliche Bieterkommunikation im Portal. Preis ist das einzige Zuschlagskriterium; Eignung und technische Mindestanforderungen werden vorgelagert geprüft. Die Ausgabe der Unterlagen wird freigegeben.')], 'Freigegeben: Mara Seidel, Vergabestelle; mitgezeichnet: Jonas Berg, 27.07.2026.')
    pdf('02_bekanntmachung.pdf',CITY,'Auftragsbekanntmachung Los 430','29.07.2026',[[
        ('1 Bekanntmachungsdaten','Bekanntmachungstyp: Wettbewerb, Bauauftrag, offenes Verfahren. Interne Bekanntmachungskennung NO-FH26-BM01, Version 1. Übermittlung am 27.07.2026 um 09:00 Uhr; Veröffentlichung am 29.07.2026. Dieser vollständige Textauszug enthält die Verfahrens- und Losangaben der veröffentlichten Bekanntmachung. Er enthält keine externe Veröffentlichungsnummer.'),
        ('2 Auftraggeber und Kommunikation','Auftraggeber ist die Stadt Northeim, Scharnhorstplatz 1, 37154 Northeim, Deutschland, kommunale Gebietskörperschaft, allgemeine öffentliche Verwaltung. Kontakt im Verfahren: Zentrale Vergabestelle, Mara Seidel, vergabe@northeim-vergabe.example. Unterlagen und elektronische Abgabe: https://portal.northeim-vergabe.example/verfahren/NO-FH26-L430. Die Unterlagen sind dort vollständig, kostenfrei und ohne Zugangsbeschränkung abrufbar. Fragen und Antworten werden über dieses Portal bearbeitet.'),
        ('3 Gegenstand und Wert','Neubau Feuerwehrhaus Northeim, Los 430 Lüftungsanlagen. Hauptklassifikation CPV 45331210: Installation von Lüftungsanlagen. Erfüllungsort Northeim, NUTS DE918, Projektstandort Am Gerätehof 4. Ausführung eines zentralen Zu-/Abluftsystems mit 6.000 m³/h, zwei Abluftventilatoren, Kanalnetz, Regelung, Brandschutzkomponenten, Messung und Inbetriebnahme. Gesamtbauwert aller Lose 8.400.000,00 EUR netto; geschätzter Wert dieses Loses 380.000,00 EUR netto. Andere Gewerke werden gesondert ausgeschrieben.')], [
        ('4 Verfahren und Ausführung','Offenes Verfahren, keine Beschleunigung. Rechtsgrundlage: Richtlinie 2014/24/EU, GWB Teil 4 und VOB/A Abschnitt 2 in der für das Verfahren maßgeblichen Fassung. Es handelt sich weder um eine Rahmenvereinbarung noch um ein dynamisches Beschaffungssystem. Keine elektronische Auktion. Keine EU-Fördermittel. Das Übereinkommen über das öffentliche Beschaffungswesen ist anwendbar. Kleine und mittlere Unternehmen können teilnehmen. Los 430 bildet innerhalb dieser Bekanntmachung den gesamten Auftragsumfang. Es bestehen keine Optionen und keine Verlängerungen.'),
        ('5 Teilnahme und Mindestanforderungen','Einzureichen sind eine Eigenerklärung über die Eintragung im Berufs- oder Handelsregister, das Nichtvorliegen zwingender oder fakultativer Ausschlussgründe sowie die ordnungsgemäße Zahlung von Steuern und Sozialbeiträgen. Wirtschaftliche Mindestanforderung: mittlerer Umsatz mit Lüftungsleistungen der letzten drei abgeschlossenen Geschäftsjahre mindestens 600.000 EUR netto. Nachweis zunächst durch Eigenerklärung. Technische Mindestanforderung: mindestens eine in den letzten fünf Jahren fertiggestellte Zu-/Abluftanlage mit mindestens 5.000 m³/h in einem Nichtwohngebäude; der Bieter muss Montage und Inbetriebnahme verantwortet haben. Benennung von Ort, Auftraggeber, Ansprechpartner, Zeitraum, Luftmenge und eigenem Umfang mit dem Angebot; Auftraggeberbestätigung auf Anforderung. Referenzen als Nachunternehmer sind zulässig. Benennung der verantwortlichen Bauleitung mit einschlägiger Berufserfahrung von mindestens drei Jahren. Keine Präqualifikation vorgeschrieben; gleichwertige Belege werden akzeptiert.'),
        ('6 Zuschlag und Nebenangebote','Preis 100 %. Maßgeblich ist die Gesamtnettosumme einschließlich erster Wartung und unbedingter Nachlässe. Qualitative Zuschlagspunkte gibt es nicht. Die technischen Mindestanforderungen ergeben sich aus dem Leistungsverzeichnis. Nebenangebote sind ausgeschlossen; ein Hauptangebot je Bieter ist zulässig. Bietergemeinschaften haften gesamtschuldnerisch und benennen einen bevollmächtigten Vertreter. Unterauftragnehmer und Umfang sind mit dem Angebot zu benennen.')], [
        ('7 Fristen und Einreichung','Angebotsfrist: Montag, 31.08.2026, 10:00 Uhr MESZ. Öffnung: 31.08.2026, 10:05 Uhr MESZ, nicht öffentlich, durch zwei Bedienstete der Vergabestelle. Sprache: Deutsch. Angebote sind ausschließlich elektronisch in Textform über das Portal einzureichen. Eine E-Mail ersetzt die Portalabgabe nicht. Die erklärende Person ist namentlich zu nennen. Angebotsbindung bis 30.10.2026, 24:00 Uhr. Rückfragen möglichst bis 20.08.2026; spätere Fragen werden bearbeitet, soweit eine rechtzeitige Information aller Bieter möglich ist. Fehlende nachforderungsfähige Unterlagen können angefordert werden; ein genereller Nachforderungsausschluss wird nicht erklärt.'),
        ('8 Vertragsbedingungen','Ausführungsbeginn 02.11.2026, Fertigstellung einschließlich Inbetriebnahme 30.04.2027. Die erste Wartung erfolgt zwölf Monate nach Abnahme. Abrechnung nach aufgemessenen Mengen und Einheitspreisen. Abschlagszahlungen monatlich nach Leistungsnachweis, Zahlung binnen 30 Kalendertagen nach Zugang einer prüfbaren Rechnung; Schlusszahlung ebenso binnen 30 Kalendertagen. Keine Vorauszahlung, keine Vertragsstrafe und keine Sicherheitsleistung. Die vollständigen besonderen Vertragsbedingungen einschließlich Arbeitsentgelt- und Unterauftragnehmerpflichten stehen in den Vergabeunterlagen.'),
        ('9 Rechtsschutz','Zuständige Nachprüfungsstelle ist die Vergabekammer Niedersachsen beim Niedersächsischen Ministerium für Wirtschaft, Verkehr und Bauen, Auf der Hude 2, 21339 Lüneburg. Telefon 04131/15-3306, E-Mail vergabekammer@mw.niedersachsen.de. Das besondere elektronische Behördenpostfach ist unter dem Suchbegriff Vergabekammer Lüneburg erreichbar. Für erkannte Verstöße gilt die Rügefrist von zehn Kalendertagen nach Paragraf 160 Absatz 3 Nummer 1 GWB. Aus Bekanntmachung oder Vergabeunterlagen erkennbare Verstöße sind spätestens bis zum Ende der Angebotsfrist zu rügen. Nach Zugang einer Nichtabhilfemitteilung dürfen bis zur Antragstellung nicht mehr als 15 Kalendertage vergehen. Die Wartepflicht nach Paragraf 134 GWB bleibt zu beachten. Weitere Auskünfte erteilt die genannte Vergabestelle.')]])
    docx('03_vergabeunterlagen.docx',CITY,'Vergabeunterlagen und Vertragsbedingungen','29.07.2026',[
        ('1 Umfang der Unterlagen','Die Ausschreibung besteht aus dieser Unterlage, dem Leistungsverzeichnis 04_lv_lueftung.xlsx und dem Angebotsformular 05_angebotsformular.docx. Verbindliche Antworten werden allen registrierten Unternehmen zeitgleich bereitgestellt. Mit dem Angebot sind das ausgefüllte Angebotsblatt einschließlich aller Einheitspreise und die Eigenerklärungen zur Eignung einzureichen. Ein inhaltsgleiches Angebotsblatt innerhalb der Kalkulationsdatei ist zulässig. Produktangaben dürfen auf einem zusätzlichen Blatt derselben Datei erfolgen. Preise sind in EUR netto mit zwei Nachkommastellen anzugeben. Mengen und Leistungsbeschreibungen dürfen nicht geändert werden.'),
        ('2 Technische Ausführungsgrundlage','Das Gebäude ist eingeschossig mit Technikraum auf der nördlichen Nebenraumspange. Zu versorgen sind Umkleiden 1.800 m³/h, Schulung 2.400 m³/h, Werkstatt 900 m³/h und Nebenräume 900 m³/h. Die Fahrzeughalle erhält die zwei getrennten Abluftventilatoren. Die Abgasabsaugung an Fahrzeugen ist nicht Bestandteil dieses Loses. Der Technikraum hat ein nutzbares Innenmaß von 6,20 x 3,80 m und 2,80 m Höhe; die Einbringöffnung beträgt 1,70 x 2,10 m. Geräte sind erforderlichenfalls in transportfähigen Segmenten einzubringen. Vor dem Gerät verbleibt ein Wartungsgang von mindestens 0,90 m. Die bauseitige Aufstellfläche ist eben und tragfähig für 1.800 kg. Der Elektroauftragnehmer stellt 400 V/32 A sowie einen Netzwerkanschluss am Schaltschrank bereit. Sämtliche Anschlüsse innerhalb des Lüftungsgewerks sind einzukalkulieren.'),
        ('3 Leistungsabgrenzung und Mengen','Das Leistungsverzeichnis umfasst 20 Positionen. Der Auftragnehmer schuldet jeweils Lieferung, Montage, Befestigung und betriebsfertigen Anschluss, soweit eine Position nicht ausdrücklich nur eine Dokumentations- oder Wartungsleistung beschreibt. Die Rechteckkanäle verlaufen oberhalb abgehängter Decken, Montagehöhe bis 3,20 m. Rundrohre verlaufen in Nebenräumen. Wanddurchbrüche werden bauseits nach rechtzeitiger Maßangabe hergestellt. Verschlüsse, eigene Abdichtungen und Brandschottungen gehören zu diesem Auftrag. Eine gesonderte Vergütung für Transport, Verpackung und übliche Hilfsmittel wird nicht vereinbart.'),
        ('PAGE',''),
        ('4 Geräteangabe und technischer Nachweis','Für Position 01.020 ist ein Hersteller und ein eindeutiger Typ mit dem Angebot zu benennen. Die Leistungsbeschreibung ist produktneutral; es gibt kein Leitfabrikat. Die Mindestwerte beziehen sich auf den Betriebspunkt 6.000 m³/h bei 600 Pa externem Druck. Für SFP gilt die Summe der elektrischen Aufnahmeleistung beider Ventilatoren bezogen auf den Luftvolumenstrom eines Luftstrangs. Herstellerdatenblatt und projektspezifische Auslegung sind auf Anforderung vorzulegen. Die Angabe eines gleichwertigen Fabrikats allein ersetzt keine eindeutige Typbezeichnung. Der Zugang für Wartung und Filtertausch muss bei den vorgegebenen Raummaßen erhalten bleiben.'),
        ('5 Eignungsunterlagen','Die in Abschnitt 5 der Bekanntmachung festgelegten Mindestanforderungen gelten unverändert. Die Referenzerklärung muss erkennen lassen, ob der Bieter für die Gesamtanlage oder nur für Teilleistungen zuständig war. Eine Tätigkeit als Nachunternehmer ist kein Ausschlussgrund. Bei getrennter Erbringung von Montage und Inbetriebnahme ist der eigene Verantwortungsumfang darzustellen. Auftraggeberbestätigungen werden zunächst nicht verlangt. Die Vergabestelle kann Belege für bereits erklärte Tatsachen nachfordern und Angaben aufklären. Sie entscheidet über eine Nachforderung im Einzelfall; eine Änderung von Leistungsinhalt oder Angebotspreisen wird damit nicht eröffnet.'),
        ('6 Preisangaben und Wertung','Jeder Einheitspreis muss die ausgeschriebene Leistung vollständig abdecken. Die Positionssumme ist Menge mal Einheitspreis, kaufmännisch auf Cent gerundet. Unbedingte Nachlässe sind als Prozentsatz auf die gesamte Nettosumme anzugeben. Skonti werden nicht gewertet. Umsatzsteuer ist gesondert mit 19 % auszuweisen. Preis ist das einzige Zuschlagskriterium. Eine rechnerische Rangfolge besagt noch nicht, dass ein Angebot technisch und hinsichtlich der Eignung berücksichtigt werden kann. Bei aufklärungsbedürftigen Preisabständen werden die betreffenden Kalkulationsbestandteile erläutert verlangt.'),
        ('PAGE',''),
        ('7 Ausführung und Abnahme','Die Ausführung beginnt am 02.11.2026 und muss am 30.04.2027 einschließlich Messung und Inbetriebnahme abgeschlossen sein. Innerhalb von zehn Arbeitstagen nach Zuschlag ist ein abgestimmter Montageablaufplan vorzulegen. Ausführungszeichnungen und Einbaumaße sind vor Bestellung der Geräte freizugeben. Eine Freigabe ersetzt nicht die Verantwortung für die angebotene Beschaffenheit. Die Abnahme wird nach Fertigstellungsanzeige gemeinsam durchgeführt und schriftlich protokolliert. Voraussetzungen sind die vollständige Funktion, die Übergabe der Bestandsunterlagen und die Einweisung. Mängelrechte richten sich nach den gesetzlichen Bestimmungen. Für dieses Vertragsverhältnis wird die VOB/B nicht zusätzlich einbezogen.'),
        ('8 Vergütung und Vertragsabwicklung','Die Vergütung bestimmt sich nach den angebotenen Einheitspreisen und den gemeinsam festgestellten tatsächlichen Mengen. Ein unbedingter Nachlass gilt für alle Positionen. Änderungen des Leistungsumfangs bedürfen vor Ausführung einer dokumentierten Anordnung der Stadt und einer gesonderten Vergütungsabstimmung; zwingende gesetzliche Ansprüche bleiben unberührt. Abschlags- und Schlussrechnungen sind mit Positionsnachweis und Aufmaß elektronisch an die Vergabestelle zu richten. Prüffähige Rechnungen werden innerhalb von 30 Kalendertagen bezahlt. Vorauszahlungen, Sicherheiten und Vertragsstrafen werden nicht vereinbart. Die erste Wartung ist bereits im angebotenen Gesamtpreis enthalten und wird nach deren Durchführung abgerechnet.'),
        ('9 Personal und Unterauftragnehmer','Der Auftragnehmer verpflichtet sich, die für seine Beschäftigten jeweils verbindlichen gesetzlichen und allgemeinverbindlichen tariflichen Mindestarbeitsbedingungen einzuhalten und diese Pflicht an eingesetzte Unterauftragnehmer weiterzugeben. Die angebotene Bauleitung muss während der Montage erreichbar sein. Ein Wechsel ist mit einer mindestens gleich erfahrenen Person rechtzeitig anzuzeigen. Der Auftragnehmer koordiniert seine Arbeiten mit den übrigen Gewerken, schützt benachbarte Bauteile und beseitigt eigene Abfälle. Für vom Auftraggeber bereitgestellte Anschlüsse und Flächen werden keine Umlagen erhoben. Bei Behinderungen ist die Bauleitung unverzüglich in Textform zu informieren.')], 'Ausgegeben durch Mara Seidel, Zentrale Vergabestelle.')
    docx('05_angebotsformular.docx',CITY,'Angebotsformular Los 430','29.07.2026',[
        ('1 Bieter und Erklärung','Firma: __________________________________________________________\nAnschrift: ________________________________________________________\nErklärende Person und Funktion: ______________________________________\nE-Mail und Telefon: _________________________________________________\nOrt und Datum: ____________________________________________________'),
        ('2 Angebot','Wir bieten die Leistungen des Loses 430 auf Grundlage der Vergabeunterlagen vom 29.07.2026 und der veröffentlichten Antworten zu den Einheitspreisen unseres beigefügten Leistungsverzeichnisses an. Die Bindung gilt bis 30.10.2026, 24:00 Uhr. Die Termine 02.11.2026 und 30.04.2027 werden eingehalten. Unsere eigenen Geschäftsbedingungen werden nicht Vertragsbestandteil.'),
        ('3 Preis und Fabrikat','Nettosumme vor Nachlass: __________________ EUR\nUnbedingter Nachlass auf alle Positionen: __________ %\nNettosumme nach Nachlass: _________________ EUR\nUmsatzsteuer 19 %: ________________________ EUR\nGesamtsumme brutto: ______________________ EUR\nHersteller und eindeutiger Typ Position 01.020: __________________________'),
        ('4 Ausführung und Beilagen','Unterauftragnehmer und Leistungsumfang oder Erklärung vollständiger Eigenleistung: ______________________________________________________\nBenannte Bauleitung: _________________________________________________\nBeigefügt sind das vollständig bepreiste Leistungsverzeichnis und die Eigenerklärungen zu Registereintragung, Ausschlussgründen, Abgaben, Umsatz, Referenz und Bauleitung. Bei Bietergemeinschaften ist die bevollmächtigte Vertretung zu benennen. Ein inhaltsgleiches Angebotsblatt in der Kalkulationsdatei ersetzt dieses Formular.')], 'Abgabe in Textform durch die oben namentlich benannte erklärende Person.')
    eml('06_bieterfrage.eml','Lena Winter <angebote@leinetal-luft.example>','Vergabestelle <vergabe@northeim-vergabe.example>','2026-08-12T10:18:00+02:00','Frage zu Referenz und Auslegungsunterlage','''Sehr geehrte Frau Seidel,

bei unserer Referenz Sportzentrum Weidenfeld hat der Generalunternehmer den Hauptvertrag gehalten. Wir haben die Lüftungsanlage montiert und die Inbetriebnahme gemeinsam mit dem Gerätehersteller durchgeführt. Genügt hierfür unsere eigene Referenzerklärung und kann die Bestätigung des Generalunternehmers nachgereicht werden?

Für Position 01.020 möchten wir einen Typ aus der Baureihe AL 6000 anbieten. Ist der vollständige projektspezifische Auslegungsbogen bereits mit dem Angebot erforderlich, oder reicht die eindeutige Typangabe mit späterem Datenblatt auf Anforderung? Wir fragen nicht nach einer Änderung der geforderten Leistungswerte.

Mit freundlichen Grüßen
Lena Winter, Geschäftsführerin
Leinetal Lufttechnik GmbH''','NO26-F006')
    eml('07_bieterantwort.eml','Vergabestelle <vergabe@northeim-vergabe.example>','Registrierte Unternehmen <teilnehmer@portal.northeim-vergabe.example>','2026-08-14T14:00:00+02:00','Verbindliche Antwort 1 an alle Unternehmen','''Sehr geehrte Damen und Herren,

zu der anonymisierten Frage vom 12.08.2026 teilen wir für alle Unternehmen mit: Eine als Nachunternehmer ausgeführte Referenz ist zulässig. Entscheidend bleibt, dass das anbietende Unternehmen Montage und Inbetriebnahme einer Anlage des geforderten Umfangs verantwortet hat. Eine Auftraggeberbestätigung ist erst auf Anforderung vorzulegen. Die eigene Erklärung muss mit dem Angebot vorliegen.

Für Position 01.020 genügt mit dem Angebot die eindeutige Hersteller- und Typbezeichnung. Herstellerdatenblatt und projektspezifische Auslegung werden auf Anforderung vorgelegt. Die Mindestwerte aus dem Leistungsverzeichnis bleiben unverändert. Eine nachträgliche Änderung des angebotenen Geräts wird mit dieser Antwort nicht zugelassen. Bei einer Typfamilie mit unterschiedlichen Baugrößen oder Ventilatorvarianten ist die konkrete angebotene Ausführung anzugeben.

Die Angebotsfrist 31.08.2026 um 10:00 Uhr MESZ bleibt unverändert. Diese Nachricht ist Bestandteil der Vergabeunterlagen.

Mara Seidel
Zentrale Vergabestelle''','NO26-F007','NO26-F006')
    for i, bidder in enumerate(BIDDERS):
        number = ['09','11','13'][i]
        if i == 0:
            ref = 'Sportzentrum Weidenfeld, Göttingen; Auftraggeber Bauverbund Weidenfeld GmbH, Projektleitung Eva Rauch, eva.rauch@bauverbund-weidenfeld.example. Ausführungszeit Mai bis Oktober 2024. Anlage mit 5.400 m³/h. Unser Unternehmen verantwortete Montage und Inbetriebnahme der Lüftungsanlage als Nachunternehmer. Der Gerätehersteller begleitete die Inbetriebnahme. Auftragssumme unseres Anteils 248.000 EUR netto. Fertigstellung am 18.10.2024. Eine Bestätigung wird auf Anforderung beschafft.'
            turnover='2023: 820.000 EUR; 2024: 910.000 EUR; 2025: 960.000 EUR.'
            lead='Tobias Renz, Meister für Installations- und Gebäudetechnik, seit 2018 in der Bauleitung, davon fünf Jahre Lüftungsanlagen.'
        elif i == 1:
            ref='Bildungszentrum Solling, Holzminden; Auftraggeber Bildungswerk Solling GmbH, Petra Linde, petra.linde@bildungswerk-solling.example. März bis November 2023, Zu-/Abluft 6.800 m³/h. Wir waren Hauptauftragnehmer für Montage, Regelung, Einregulierung und Inbetriebnahme. Fertigstellung 24.11.2023, eigener Auftragswert 412.000 EUR netto. Zweite Referenz: Verwaltungshaus Uferhof, Höxter, Uferhof Immobilien GmbH, Ralf Meier, ralf.meier@uferhof.example; Juni bis Dezember 2025, 5.200 m³/h, Montage und Inbetriebnahme, 338.000 EUR netto, fertiggestellt 12.12.2025.'
            turnover='2023: 1.460.000 EUR; 2024: 1.620.000 EUR; 2025: 1.710.000 EUR.'
            lead='Annika Friese, staatlich geprüfte Technikerin Versorgungstechnik, seit 2016 Bauleiterin für Lüftungsanlagen.'
        else:
            ref='Logistikverwaltung Harztor, Seesen; Auftraggeber Harztor Gewerbebau GmbH, Jens Voss, jens.voss@harztor-gewerbebau.example. Februar bis September 2024, Zu-/Abluft 7.200 m³/h. Wir führten Montage und Inbetriebnahme einschließlich Regelungsprüfung in eigener Verantwortung aus. Fertigstellung 20.09.2024, Auftragswert 468.000 EUR netto. Unsere Bauleitung führte die Mess- und Funktionsprotokolle.'
            turnover='2023: 1.120.000 EUR; 2024: 1.340.000 EUR; 2025: 1.280.000 EUR.'
            lead='Daniel Hohmann, Meister für Lüftungsbau, seit 2017 verantwortlich für Montage und Inbetriebnahme.'
        pdf(f'{number}_eignung_{bidder["short"]}.pdf',bidder['name'],'Eigenerklärung zur Eignung',bidder['date'],[[
            ('1 Unternehmen und Register',f'{bidder["name"]}, {bidder["address"]}. Wir sind im Handelsregister als Gesellschaft mit beschränkter Haftung eingetragen und zur Ausführung der angebotenen Leistungen berechtigt. Die Geschäftsführung vertritt das Unternehmen. Ein Registerauszug und eine aktuelle Bescheinigung der zuständigen Handwerkskammer werden auf Anforderung vorgelegt. Erklärende Person: {bidder["person"]}, Geschäftsführung.'),
            ('2 Ausschlussgründe und Abgaben','Wir erklären, dass nach Prüfung unserer Geschäftsunterlagen keine Umstände bestehen, die einen zwingenden oder fakultativen Ausschluss aus diesem Verfahren begründen. Insbesondere sind uns keine einschlägigen rechtskräftigen Verurteilungen unserer vertretungsberechtigten Personen, keine Insolvenzverfahren und keine schweren beruflichen Verfehlungen bekannt. Steuern, Abgaben und Sozialversicherungsbeiträge sind ordnungsgemäß entrichtet. Es bestehen keine Rückstände aus fälligen, unbestrittenen Verpflichtungen. Änderungen werden unverzüglich mitgeteilt.'),
            ('3 Wirtschaftliche Leistungsfähigkeit',f'Umsatz ausschließlich mit Lüftungsleistungen, jeweils netto: {turnover} Die Angaben beziehen sich auf abgeschlossene Geschäftsjahre und nicht auf die gesamte Unternehmensgruppe. Jahresabschlussauszüge und Abgabenbescheinigungen werden auf Anforderung bereitgestellt. Eine Betriebshaftpflicht mit Deckung von 3 Mio. EUR für Personen- und Sachschäden wird während der Ausführung unterhalten.')],[
            ('4 Referenz über eine fertiggestellte Anlage',ref),
            ('5 Personal und Ausführung',f'Verantwortliche Bauleitung: {lead} Wir verfügen über sechs eigene Monteure sowie die erforderlichen Messgeräte und Montagewerkzeuge. Die ausgeschriebenen Arbeiten werden in Eigenleistung ausgeführt. Die Lieferung und produktbezogene Unterstützung durch den Gerätehersteller ist kein Nachunternehmerauftrag für die Montage. Eine Eignungsleihe oder Bietergemeinschaft liegt nicht vor.'),
            ('6 Bestätigung',f'Wir bestätigen die Richtigkeit dieser Angaben und gestatten der Vergabestelle die Kontaktaufnahme mit den benannten Referenzgebern zur Prüfung des angegebenen Leistungsumfangs.\n{bidder["person"]}, Geschäftsführung\n{bidder["date"]}')]])
    docx('14_oeffnungsniederschrift.docx',CITY,'Niederschrift über die Angebotsöffnung','31.08.2026',[
        ('1 Öffnung','Die Angebotsfrist endete am 31.08.2026 um 10:00:00 Uhr MESZ. Mara Seidel und Jonas Berg öffneten die verschlüsselten Angebote am selben Tag von 10:05 bis 10:23 Uhr. Die Zugriffssperre war zuvor aktiv. Es liegen drei fristgerechte Hauptangebote und keine Nebenangebote vor. Keine Rücknahme und kein verspäteter Eingang wurden angezeigt.'),
        ('2 Eingegangene Angebote',f'NO26-E001: {ALPHA}, Eingang 31.08.2026, 09:54:12 Uhr; Nettosumme {a} EUR, Nachlass 0 %. NO26-E002: {BETA}, Eingang 31.08.2026, 09:31:40 Uhr; Nettosumme {b} EUR, Nachlass 0 %. NO26-E003: {GAMMA}, Eingang 31.08.2026, 09:48:05 Uhr; Nettosumme {c} EUR, Nachlass 0 %. Die Beträge wurden aus den Angebotsblättern abgelesen; eine fachliche Entscheidung wird hier nicht getroffen.'),
        ('3 Dokumente und Auffälligkeiten','Jedes Paket enthält ein XLSX-Angebot und eine PDF-Eigenerklärung. Alle erklärenden Personen sind benannt. Bei Leinetal ist als Gerät Auenluft AL 6000 K eingetragen. Die Referenz beschreibt eine Tätigkeit als Nachunternehmer. Die Herstellerunterlagen wurden entsprechend der Ausschreibung noch nicht beigefügt. Alle drei Dateien enthalten Preise für sämtliche Positionen. Die Unterlagen werden unverändert für die fachliche Prüfung gesichert.')], 'Mara Seidel und Jonas Berg, abgeschlossen am 31.08.2026 um 10:23 Uhr.')
    screens()
    docx('18_nachforderung_aufklaerung.docx',CITY,'Anforderung von Nachweisen und Aufklärung','02.09.2026',[
        ('',f'An {ALPHA}, Frau Lena Winter, über das Vergabeportal. Bezug: Ihr Angebot, Eingang NO26-E001. Bitte beantworten Sie die folgenden Punkte bis Mittwoch, 09.09.2026, 12:00 Uhr MESZ, über das Portal. Eingang und Abruf sind dort dokumentiert.'),
        ('1 Referenzbestätigung','Bitte legen Sie für das Sportzentrum Weidenfeld eine Bestätigung des benannten Auftraggebers vor, aus der Fertigstellung, Luftmenge und der von Ihnen verantwortete Umfang von Montage und Inbetriebnahme hervorgehen. Wir fordern einen Beleg für die mit dem Angebot erklärte Referenz an, keine neue Referenz.'),
        ('2 Geräteausführung','Bitte legen Sie das Herstellerdatenblatt und die projektspezifische Auslegung für den angebotenen Typ Auenluft AL 6000 K vor. Insbesondere sind SFP, Wärmerückgewinnung und Gehäuseabmessungen am ausgeschriebenen Betriebspunkt zu belegen. Bitte erläutern Sie, ob der Typ Zusatzmodule oder verschiedene Ventilatorausführungen umfasst und welche Ausführung Ihrer Preisbildung zugrunde lag.'),
        ('3 Preisaufklärung','Der Einheitspreis der Position 01.020 beträgt bei Ihnen 78.000,00 EUR netto. Er liegt unter den weiteren Angebotspreisen von 102.000,00 EUR und 107.500,00 EUR. Bitte erläutern Sie die Kalkulation dieser Position sowie die vollständige Einbeziehung der Regelung und Inbetriebnahme in die dafür vorgesehenen Positionen. Eine Änderung von Einheitspreisen oder angebotener Leistung wird nicht angefordert.'),
        ('4 Rückmeldung','Sollte ein Beleg nicht rechtzeitig vorliegen, erläutern Sie dies unverzüglich konkret. Die Vergabestelle prüft anschließend, ob und in welchem Umfang die Unterlagen verwertet werden können. Bis dahin ist keine Entscheidung über die Berücksichtigung Ihres Angebots getroffen.')], 'Mara Seidel, Zentrale Vergabestelle; versandt 02.09.2026, 09:00 Uhr MESZ, NO26-N018.')
    docx('19_antwort.docx',ALPHA,'Antwort auf die Anforderung vom 2 September','09.09.2026',[
        ('','An die Stadt Northeim, Zentrale Vergabestelle. Wir reichen die Auftraggeberbestätigung vom 07.09.2026 und das technische Datenblatt vom 08.09.2026 ein. Die angebotenen Preise und Leistungen bleiben unverändert.'),
        ('1 Referenz','Unsere Aussage zur Verantwortung für die Inbetriebnahme bezog sich auf die Gesamtkoordination unseres Gewerks einschließlich der Messungen und der Übergabe an den Generalunternehmer. Der Hersteller stellte die Regelparameter ein. Wir veranlassten und begleiteten diese Tätigkeit und stellten die betriebsbereite Anlage zur Abnahme vor. Die beigefügte Bestätigung beschreibt die Arbeiten nach der beim Generalunternehmer vorhandenen Abrechnung. Ein von beiden Unternehmen unterzeichnetes Inbetriebnahmeprotokoll haben wir derzeit nicht auffinden können.'),
        ('2 Typbezeichnung','AL 6000 K ist die in unserem Angebot benannte Kompaktbaureihe. Angeboten war nach unserer Kalkulation die werkseitige EC-Ausführung mit Ventilatorsatz E2. Das Kürzel E2 wird in der Herstellerbestellung als Konfiguration geführt, nicht als eigener Gerätetyp. Der Auslegungsbogen zeigt diese Ausführung. Wir ändern das Gerät nicht. Eine vor dem 31.08.2026 datierte Herstellerbestellbestätigung existiert nicht, weil wir erst nach Zuschlag bestellen. Unser Kalkulationsgespräch mit dem Hersteller fand am 26.08.2026 statt; eine gesonderte schriftliche Auslegung wurde damals nicht erstellt.'),
        ('3 Kalkulation Position 01.020','Der Einheitspreis von 78.000,00 EUR netto setzt sich aus 61.000,00 EUR Geräteeinkauf einschließlich Ventilatorsatz E2, 3.200,00 EUR Transport und Hebevorgang, 7.600,00 EUR Montagepersonal, 1.200,00 EUR Anschlussmaterial sowie 5.000,00 EUR Gemeinkosten und Ergebnis zusammen. Den Einkaufspreis trägt unser Jahreskonditionenrahmen. Wir tragen das Risiko einer Abweichung vom kalkulierten Einkaufspreis. Es sind keine Fördermittel eingerechnet.'),
        ('4 Abgrenzung','Regelung, Schaltschrank und Feldgeräte sind mit 24.500,00 EUR in Position 01.120 enthalten. Die Inbetriebnahmeleistungen einschließlich Herstellerunterstützung sind mit 9.200,00 EUR in Position 01.160 enthalten. Die Positionspreise werden nicht untereinander verschoben. Für die vollständige Leistung stehen wir zum eingereichten Gesamtpreis ein.')], 'Lena Winter, Geschäftsführerin; übermittelt 09.09.2026 um 11:41:02 Uhr MESZ, NO26-A019.')
    pdf('20_referenz.pdf','Bauverbund Weidenfeld GmbH','Bestätigung der ausgeführten Leistungen','07.09.2026',[[
        ('','An Leinetal Lufttechnik GmbH, Frau Winter. Auf Ihre Bitte bestätigen wir die bei uns dokumentierten Leistungen am Sportzentrum Weidenfeld in Göttingen.'),
        ('1 Beauftragung und Fertigstellung','Wir beauftragten Leinetal mit Lieferung und Montage der Zu-/Abluftanlage mit einer Nennluftmenge von 5.400 m³/h. Die Arbeiten wurden zwischen Mai und Oktober 2024 erbracht. Unser Auftrag an Leinetal hatte einen abgerechneten Wert von 248.000,00 EUR netto. Die Anlage wurde am 18.10.2024 als betriebsbereit an uns übergeben. Die bei Übergabe vermerkten Beschriftungsmängel wurden im November 2024 beseitigt.'),
        ('2 Inbetriebnahme','Leinetal führte die Volumenstrommessung und die mechanische Einregulierung aus. Die erstmalige Parametrierung der Gerätesteuerung und die Freigabe der Sicherheitskette wurden durch einen Servicetechniker des Geräteherstellers vorgenommen. Dieser Termin war durch Leinetal organisiert; Herr Renz war anwesend. Wer im Innenverhältnis die Herstellerleistung beauftragt und bezahlt hat, lässt sich aus unserer Projektakte nicht erkennen. Unsere technische Abnahme richtete sich an Leinetal als unseren Vertragspartner.'),
        ('3 Aussageumfang','Diese Bestätigung beschreibt die von uns beobachteten und abgerechneten Leistungen. Eine Bewertung von Eignungsvoraussetzungen in anderen Vergabeverfahren geben wir nicht ab. Für Rückfragen steht unsere Projektleitung zur Verfügung.\nEva Rauch, Projektleitung\neva.rauch@bauverbund-weidenfeld.example')]])
    pdf('21_geraet.pdf','Auenluft Gerätebau GmbH','Projektauslegung AL 6000 K','08.09.2026',[[
        ('1 Auftrag und Ausführung','Auslegung NO-AL-260908-01 für Leinetal Lufttechnik GmbH, Projekt Feuerwehrhaus Northeim. Ansprechpartner: Felix Ahrens, technische Auslegung. Gerätebaureihe AL 6000 K, Gehäuseausführung K, Ventilatorsatz E2, Gegenstrom-Wärmerückgewinnung. Diese Auslegung wurde am 08.09.2026 erstellt. Sie ist keine Bestellung und keine Lieferfreigabe.'),
        ('2 Betriebspunkt','Zuluft 6.000 m³/h; Abluft 6.000 m³/h; externer Druck jeweils 600 Pa. Elektrische Gesamtaufnahme beider Ventilatoren 2,80 kW, entsprechend SFP 1,68 kW/(m³/s). Trockener Temperaturübertragungsgrad der Wärmerückgewinnung 82 %. Gehäuse 3.150 x 1.580 x 1.760 mm, Gewicht 1.540 kg. Geteilte Lieferung in drei Segmenten; größtes Segment 1.580 x 1.720 mm im Einbringquerschnitt. Servicezugang auf einer Längsseite, erforderlicher freier Gang 0,90 m. Schnittstelle BACnet/IP ist mit dem angebotenen Regelpaket verfügbar.'),
        ('3 Auswahlstand','Die Baureihe wird mit den Ventilatorsätzen E1 und E2 angeboten. Bei E1 liegt die elektrische Gesamtaufnahme am genannten Betriebspunkt bei 3,15 kW, entsprechend SFP 1,89 kW/(m³/s). Der äußere Gerätetyp wird in beiden Fällen als AL 6000 K bezeichnet. Der Ventilatorsatz muss in der Bestellung gesondert bestimmt werden. Beide Sätze sind werkseitige Konfigurationen; ein Wechsel auf der Baustelle ist nicht vorgesehen. Der Auslegung E2 liegt die telefonische Angabe von Frau Winter vom 07.09.2026 zugrunde. Die Gesprächsnotiz vom 26.08.2026 nennt den Betriebspunkt und die Baureihe, aber keinen Ventilatorsatz.'),
        ('4 Erklärung','Die genannten Werte gelten für die hier ausgewählte Konfiguration. Wir können aus unseren Unterlagen nicht bestätigen, welche Konfiguration der Kalkulation des Bieters vor Angebotsabgabe zugrunde lag. Die aktuelle technische Lieferbarkeit der Ausführung E2 wird bestätigt.\nFelix Ahrens, technische Auslegung')]])
    (CASE/'22_chat_fachpruefung.txt').write_text('''Projektchat Hochbau / Vergabestelle
Exportiert von Nora Brandt am 10.09.2026, 16:40 Uhr MESZ
Vorgang NO-FH26-L430

10.09.2026 09:12 Nora Brandt: Mit E2 passen die Auslegungswerte und die Einbringmaße. E1 liegt beim SFP darüber. Im Angebotsblatt steht nur AL 6000 K.
10.09.2026 09:18 Mara Seidel: Ist E2 ein anderes Gerät oder nur die Konfiguration? Ich brauche eine Aussage zur technischen Identität, nicht zur Wertung.
10.09.2026 09:24 Nora Brandt: Die Gehäusebaureihe ist gleich. Die Ventilatorsätze sind verschieden. Ich kann aus den Unterlagen nicht feststellen, was am 31. August kalkuliert war.
10.09.2026 10:03 Jonas Berg: Bei der Referenz bestätigt Rauch die Übergabe an Leinetal. Sie sagt gleichzeitig, die Steuerung kam vom Hersteller. Unsere Antwort an alle ließ Herstellerbeteiligung nicht ausdrücklich außen vor.
10.09.2026 10:11 Mara Seidel: Ich lese unsere Mindestanforderung als Verantwortung für Montage und Inbetriebnahme. Die Bestätigung muss im Zusammenhang mit der Eigenerklärung geprüft werden.
10.09.2026 15:32 Nora Brandt: Die Kostenaufteilung für 01.020 ergibt 78.000 EUR. Regelung und Inbetriebnahme sind gesondert kalkuliert. Keine offensichtliche Doppelerfassung. Einkaufskonditionen selbst habe ich nicht verifiziert.
10.09.2026 16:02 Jonas Berg: Bitte beide Punkte im Vermerk sichtbar lassen. Ein Gesprächsvermerk ist kein Nachweis des Angebotsinhalts.
''',encoding='utf-8')
    docx('24_wertungsvermerk.docx',CITY,'Wertungsvermerk und Vergabevorschlag','14.09.2026',[
        ('1 Form und Preis',f'Drei fristgerechte und namentlich erklärte Hauptangebote liegen vor. Die Rechenprüfung im Preisspiegel ergibt {ALPHA} {a} EUR, {BETA} {b} EUR und {GAMMA} {c} EUR netto. Alle 20 Positionen sind bepreist. Keine Nachlässe oder Nebenbedingungen sind erklärt. Die Preisaufklärung von Leinetal summiert sich auf den angebotenen Einheitspreis. Die Kalkulation der Regelung und Inbetriebnahme ist gesondert erläutert. Ein eigenständiger Ausschluss wegen nicht tragfähiger Preise wird auf dieser Grundlage nicht vorgeschlagen.'),
        ('2 Eignung','Weserklima und Harzraum erfüllen nach den Eigenerklärungen die veröffentlichten Umsatz-, Personal- und Referenzanforderungen. Zu Leinetal liegen Eigenerklärung, Aufklärungsantwort und Bestätigung des Generalunternehmers vor. Die Bestätigung dokumentiert Montage, mechanische Einregulierung und Übergabe durch Leinetal. Parametrierung und Freigabe der Sicherheitskette wurden vom Hersteller erbracht. Mara Seidel hält den Nachweis der geforderten Verantwortung für die gesamte Inbetriebnahme damit nicht für ausreichend. Jonas Berg weist darauf hin, dass die technische Abnahme gegenüber Leinetal erfolgte und die Ausschreibung keine vollständige Eigenleistung bei der Inbetriebnahme verlangte. Ein unterschriebenes Inbetriebnahmeprotokoll liegt nicht vor.'),
        ('3 Technischer Angebotsinhalt','Weserklima bietet Flusswind FW 6200 mit im Angebotsblatt angegebenen Werten von 1,74 kW/(m³/s), 81 % Wärmerückgewinnung und Gehäuse 3.100 x 1.550 x 1.750 mm an. Harzraum bietet Nordvent NV 6000 mit 1,71 kW/(m³/s), 83 % und 3.180 x 1.590 x 1.780 mm an. Die selbst erklärten Werte erfüllen die Mindestanforderungen. Bei Leinetal weist die nachgereichte Auslegung E2 Werte innerhalb der Anforderungen aus. Das eingereichte Angebotsblatt benennt jedoch keinen Ventilatorsatz. Nach Herstellerangabe ist auch E1 unter derselben Baureihenbezeichnung lieferbar. Eine vor Angebotsabgabe dokumentierte Auswahl von E2 liegt nicht vor.'),
        ('PAGE',''),
        ('4 Bewertung und Entscheidungsvorschlag','Mara Seidel schlägt vor, das Angebot von Leinetal nicht zu berücksichtigen. Ausschlaggebend sind aus ihrer Sicht die nicht hinreichend belegte Referenzverantwortung und die erst in der Aufklärung konkret benannte Geräteausführung. Die Vergabestelle behandelt die Auslegung nicht als bloße Bestätigung eines bereits eindeutig festgelegten Angebotsinhalts. Diese Einschätzung wird von Jonas Berg hinsichtlich der Geräteangabe mitgetragen; zur Referenz verbleibt der in Abschnitt 2 dokumentierte Vorbehalt. Nora Brandt bestätigt nur die technische Eignung der Konfiguration E2, nicht deren Festlegung bei Angebotsabgabe.'),
        ('5 Weiteres Vorgehen','Unter den nach dieser Prüfung berücksichtigten Angeboten hat Weserklima den niedrigsten Preis. Die Vergabe an Weserklima wird vorgeschlagen. Die Informationen an Leinetal und Harzraum werden am 14.09.2026 elektronisch versandt. Als frühester Vertragsschluss wird 25.09.2026 benannt. Ein Zuschlag wird heute nicht erteilt. Vor einer Zuschlagserteilung sind die abschließende Registerprüfung und etwaige Rechtsschutzmitteilungen zu beachten.')], 'Entscheidung: Vorschlag freigegeben. Mara Seidel, 14.09.2026; Mitzeichnung Jonas Berg mit dokumentiertem Vorbehalt zu Abschnitt 2.')
    for n, bidder, reason in [
        ('25',ALPHA,'Ihr Angebot wird nicht berücksichtigt, weil wir den Nachweis der geforderten Verantwortung für Montage und Inbetriebnahme der Referenzanlage auch nach Vorlage der Auftraggeberbestätigung nicht für ausreichend halten. Die Bestätigung weist wesentliche Inbetriebnahmetätigkeiten dem Gerätehersteller zu. Zudem ist aus Ihrer ursprünglichen Geräteangabe AL 6000 K die Konfiguration E2 nicht eindeutig erkennbar. Erst die Aufklärungsantwort benennt diese Ausführung; die ebenfalls unter derselben Baureihe verfügbare E1-Ausführung erfüllt den SFP-Mindestwert nicht. Wir behandeln die spätere Konkretisierung deshalb nicht als Bestätigung eines bereits eindeutig beschriebenen Angebots. Der aufgeklärte Preisabstand ist kein weiterer Nichtberücksichtigungsgrund.'),
        ('26',GAMMA,f'Ihr Angebot erfüllt nach unserer Prüfung die Teilnahme- und technischen Mindestanforderungen. Sein Nettopreis von {c} EUR liegt über dem Nettopreis des nach unserer Prüfung günstigsten berücksichtigten Angebots von Weserklima mit {b} EUR. Bei dem ausschließlich maßgeblichen Zuschlagskriterium Preis kann Ihr Angebot daher nicht für den Zuschlag vorgesehen werden.')]:
        pdf(f'{n}_information_{"leinetal" if n=="25" else "harzraum"}.pdf',CITY,'Information über die beabsichtigte Vergabe','14.09.2026',[[
            ('',f'An {bidder}, über das Vergabeportal. Information nach Paragraf 134 GWB zum Verfahren {AKTE}.'),
            ('1 Beabsichtigter Zuschlag',f'Die Stadt beabsichtigt, den Zuschlag auf das Angebot der {BETA}, Gewerbering 26, 37603 Holzminden, zu erteilen.'),
            ('2 Gründe Ihrer Nichtberücksichtigung',reason),
            ('3 Frühester Vertragsschluss','Diese Information wird am 14.09.2026 um 15:00 Uhr MESZ elektronisch abgesendet. Der früheste Zeitpunkt des Vertragsschlusses ist der 25.09.2026, 00:00 Uhr MESZ. Die zehn Kalendertage der Wartefrist beginnen am 15.09.2026 und enden mit Ablauf des 24.09.2026. Ein Vertrag ist noch nicht geschlossen.'),
            ('4 Kontakt','Bitte richten Sie Rückfragen unter Angabe des Verfahrenszeichens an die Zentrale Vergabestelle.\nMara Seidel\nZentrale Vergabestelle')]])
    eml('27_ruege.eml','Lena Winter <angebote@leinetal-luft.example>','Vergabestelle <vergabe@northeim-vergabe.example>','2026-09-16T09:20:00+02:00','Rüge der Nichtberücksichtigung','''Sehr geehrte Frau Seidel,

wir rügen die mit Ihrer Information vom 14.09.2026 bekanntgegebene Nichtberücksichtigung unseres Angebots. Erst diese Information hat uns die konkreten Ausschlussgründe mitgeteilt. Wir verlangen, die Entscheidung aufzuheben und unser Angebot auf Grundlage der veröffentlichten Anforderungen erneut zu prüfen. Bitte helfen Sie bis 18.09.2026, 12:00 Uhr, ab oder erklären Sie, ob Sie an der Entscheidung festhalten.

Unsere Referenz entspricht der verlangten Verantwortung für Montage und Inbetriebnahme. Wir waren Vertragspartner für die fertige Anlage. Die Beteiligung des Herstellers an der Parametrierung nimmt uns diese Verantwortung nicht. Ihre Antwort vom 14.08.2026 bestätigte ausdrücklich, dass Referenzen als Nachunternehmer zulässig sind. Die Auftraggeberbestätigung ist ein Beleg für die ursprüngliche Erklärung, kein Austausch der Referenz.

AL 6000 K war unser fest angebotenes Gerät. Wir haben die Mindestwerte vorbehaltlos zugesagt und E2 kalkuliert. Die Auslegung vom 08.09.2026 erläutert die werkseitige Konfiguration. Ein anderes Gehäuse, ein anderer Preis oder eine zusätzliche Leistung wird nicht angeboten. Aus Ihrer Antwort vom 14.08.2026 ergab sich, dass Datenblatt und Auslegung erst auf Anforderung verlangt werden. Die nachträgliche Vorlage darf nicht allein wegen ihres Datums als Angebotsänderung behandelt werden.

Unser Angebot ist nach den uns mitgeteilten Preisen günstiger. Durch den Ausschluss verlieren wir die konkrete Aussicht auf den Auftrag. Wir behalten uns einen Antrag bei der Vergabekammer vor. Bitte bestätigen Sie den Eingang.

Lena Winter
Geschäftsführerin, Leinetal Lufttechnik GmbH''','NO26-R027')
    docx('28_nichtabhilfe.docx',CITY,'Stellungnahme zur Rüge','18.09.2026',[
        ('','An Leinetal Lufttechnik GmbH, Frau Lena Winter. Ihre Rüge vom 16.09.2026 ist am selben Tag um 09:20:04 Uhr bei uns eingegangen. Wir helfen der Rüge nicht ab.'),
        ('1 Referenzverantwortung','Unsere Anforderung betrifft nicht den Status als Haupt- oder Nachunternehmer. Maßgeblich ist die Verantwortung für Montage und Inbetriebnahme. Aus der Bestätigung des Generalunternehmers ergibt sich aus unserer Sicht nicht eindeutig, dass Sie auch die durch den Gerätehersteller durchgeführte Parametrierung und Sicherheitsfreigabe verantwortet haben. Wir berücksichtigen Ihre Darstellung zur Koordination, sehen den geforderten Nachweis aber weiterhin nicht als geführt an.'),
        ('2 Geräteauswahl','Dass der technische Nachweis erst auf Anforderung einzureichen war, ist unstreitig. Bereits mit dem Angebot musste jedoch eine eindeutige Ausführung benannt sein. Die Auslegung bestätigt, dass die Baureihe mehrere Ventilatorsätze umfasst. Nach der Herstellerunterlage beruht die Auswahl E2 auf Ihrer Mitteilung vom 07.09.2026. Eine Unterlage, die E2 der Angebotsfassung vom 31.08.2026 zuordnet, ist nicht beigefügt. Ihre Zusage der Leistungswerte ist berücksichtigt, beseitigt nach unserer Auffassung die Unklarheit der Fabrikatsangabe aber nicht.'),
        ('3 Verfahrensstand','Wir halten daher an der angekündigten Nichtberücksichtigung und am Vergabevorschlag zugunsten von Weserklima fest. Die Preisaufklärung wird nicht als weiterer Ausschlussgrund verwendet. Die Frist aus der Information vom 14.09.2026 wird durch dieses Schreiben nicht neu in Gang gesetzt. Frühester Vertragsschluss bleibt der 25.09.2026. Bis heute ist kein Zuschlag erfolgt.')], 'Mara Seidel; elektronisch versandt und im Bieterpostfach zugestellt am 18.09.2026 um 11:00:00 Uhr MESZ.')
    docx('29_nachpruefungsantrag.docx','Leinetal Lufttechnik GmbH | Geschäftsführung','Antrag auf Nachprüfung','24.09.2026',[
        ('','An die Vergabekammer Niedersachsen beim Niedersächsischen Ministerium für Wirtschaft, Verkehr und Bauen, Auf der Hude 2, 21339 Lüneburg. Elektronische Einreichung durch die Geschäftsführung. Antragstellerin: Leinetal Lufttechnik GmbH, Am Werkbogen 8, 37154 Northeim, vertreten durch Geschäftsführerin Lena Winter. Antragsgegnerin: Stadt Northeim, Scharnhorstplatz 1, 37154 Northeim. Vorgesehenes Zuschlagsunternehmen: Weserklima Anlagenbau GmbH, Gewerbering 26, 37603 Holzminden. Unser Zeichen LT-NO26-NP01.'),
        ('1 Anträge','Wir beantragen, der Antragsgegnerin aufzugeben, die Entscheidung über unsere Nichtberücksichtigung im Verfahren NO-FH26-L430 aufzuheben und die Prüfung und Wertung unseres Angebots unter Beachtung der Rechtsauffassung der Vergabekammer zu wiederholen. Ferner beantragen wir Einsicht in die Vergabeakte, soweit keine berechtigten Geheimhaltungsinteressen entgegenstehen, insbesondere in die Dokumentation zu den Gründen unserer Nichtberücksichtigung. Wir bitten um unverzügliche Information der Antragsgegnerin über diesen Antrag vor dem angekündigten Vertragsschluss. Die Kosten des Verfahrens soll die Antragsgegnerin tragen.'),
        ('2 Interesse und drohender Schaden',f'Wir haben fristgerecht ein Angebot über {a} EUR netto eingereicht und wollen den Auftrag ausführen. Nach den uns mitgeteilten Vergleichspreisen liegt unser Preis unter dem Preis von Weserklima. Unser Ausschluss nimmt uns daher die konkrete Zuschlagschance. Die Auftraggeberin hat in der Bekanntmachung 8.400.000,00 EUR netto für das Gesamtbauvorhaben angegeben. Der Streit betrifft dessen Los 430, nicht eine isolierte Kleinbeschaffung.'),
        ('3 Sachverhalt und Rüge','Die Bekanntmachung datiert vom 29.07.2026. Mit Antwort vom 14.08.2026 ließ die Stadt die Referenz als Nachunternehmer zu und verlangte Auslegungsunterlagen erst auf Anforderung. Wir reichten das Angebot am 31.08.2026 um 09:54:12 Uhr ein. Auf die Anforderung vom 02.09.2026 antworteten wir am 09.09.2026 innerhalb der gesetzten Frist. Erst durch die Information vom 14.09.2026 erfuhren wir die Gründe der Nichtberücksichtigung. Wir rügten am 16.09.2026; die Nichtabhilfe ging am 18.09.2026 um 11:00 Uhr ein. Der heutige Antrag stützt sich auf Paragrafen 160 und 161 GWB und richtet sich gegen die konkrete Behandlung unseres Angebots. Er begehrt keine nachträgliche Änderung der bekanntgegebenen Mindestwerte.'),
        ('PAGE',''),
        ('4 Beanstandete Bewertung der Referenz','Die Antragsgegnerin verengt die veröffentlichte Anforderung nach unserer Auffassung auf eigenhändig ausgeführte Parametrierung. Verlangt war Verantwortung für die Inbetriebnahme. Als Vertragspartner schuldeten wir die betriebsbereite Anlage und organisierten die Herstellerunterstützung. Der Generalunternehmer nahm die Anlage gegenüber uns ab. Das nachgereichte Schreiben belegt genau diese schon ursprünglich benannte Referenz. Ob die Herstellerrechnung noch vorhanden ist, ändert nichts an unserem übernommenen Leistungsumfang. Wir machen eine Verletzung unseres Anspruchs auf Beachtung der Vergabebestimmungen nach Paragraf 97 Absatz 6 GWB und der veröffentlichten Eignungskriterien geltend.'),
        ('5 Beanstandete Behandlung der Auslegung','Wir haben das Gerät AL 6000 K mit der Zusage aller Mindestwerte angeboten. E2 ist eine Konfiguration dieses Typs. Die Antragsgegnerin behandelt eine auf ihr Verlangen erstellte Auslegung als Austausch des Angebotsinhalts, ohne eine abweichende ursprüngliche Ausführung festzustellen. Richtig ist, dass unsere Angebotsdatei E2 nicht gesondert nennt und die Herstellerunterlage keine schriftliche Auswahl vor Fristablauf bestätigt. Daraus folgt aus unserer Sicht aber nicht, dass wir E1 angeboten haben oder nachträglich ein anderes Gerät auswählen wollten. Diese Umstände sind im Zusammenhang mit unserer unveränderten Leistungserklärung und der zugelassenen späteren Nachweisführung zu würdigen.'),
        ('6 Beweismittel und Anlagen','Beigefügt werden die vollständigen Dateien 02_bekanntmachung.pdf, 03_vergabeunterlagen.docx, 04_lv_lueftung.xlsx, 06_bieterfrage.eml, 07_bieterantwort.eml, 08_angebot_leinetal.xlsx, 09_eignung_leinetal.pdf, 15_portal_bieter.png, 18_nachforderung_aufklaerung.docx, 19_antwort.docx, 20_referenz.pdf, 21_geraet.pdf, 25_information_leinetal.pdf, 27_ruege.eml und 28_nichtabhilfe.docx. Zum Referenzumfang benennen wir außerdem Eva Rauch als Auskunftsperson. Wir bitten die Kammer, die Wertungsunterlagen der Stadt beizuziehen. Die übrigen Angebote und der interne Wertungsvermerk sind uns nicht bekannt.'),
        ('7 Vertraulichkeit und Erreichbarkeit','Die Einzelpreise und die Kostenaufteilung in unserer Antwort enthalten Betriebs- und Geschäftsgeheimnisse. Vor Weitergabe dieser Angaben bitten wir um Gelegenheit zur Kennzeichnung einer geschwärzten Fassung. Eine pauschale Geheimhaltung des Verfahrensvortrags beantragen wir nicht. Zustellungen können elektronisch an die Geschäftsführung erfolgen. Ansprechpartnerin ist Lena Winter, angebote@leinetal-luft.example. Der Antrag wird von der Geschäftsführerin selbst erklärt; ein anwaltlicher Bevollmächtigter ist nicht bestellt.')], 'Lena Winter, Geschäftsführerin der Leinetal Lufttechnik GmbH, 24.09.2026.')
    # Hashes des tatsächlich beigefügten Pakets werden im Versandbeleg nach dem XLSX-Bau ergänzt.
    eml('31_vk_mitteilung.eml','Vergabekammer <geschaeftsstelle@vk-niedersachsen.example>','Vergabestelle <vergabe@northeim-vergabe.example>','2026-09-25T09:10:00+02:00','Mitteilung über den Nachprüfungsantrag','''Stadt Northeim
Zentrale Vergabestelle

Unter dem Geschäftszeichen VK-NO26-L430-N1 wird der am 24.09.2026 um 11:24:08 Uhr elektronisch eingegangene Antrag der Leinetal Lufttechnik GmbH betreffend Ihr Vergabeverfahren NO-FH26-L430 geführt. Der Antrag und seine fünfzehn Anlagen stehen in der zugehörigen gesicherten Behördennachricht VK26-OUT031 zur Verfügung. Diese E-Mail gibt die gleichzeitig übermittelte Mitteilung wieder.

Hiermit informiere ich Sie als hauptamtliche Beisitzerin über den Antrag auf Nachprüfung. Das Zuschlagsverbot nach Paragraf 169 Absatz 1 GWB ist zu beachten. Eine Entscheidung über den Antrag oder eine Gestattung des Zuschlags liegt nicht vor.

Bitte legen Sie die vollständige Vergabeakte einschließlich elektronischer Eingangs- und Versandprotokolle vor und nehmen Sie bis 30.09.2026, 12:00 Uhr MESZ, Stellung. Kennzeichnen Sie etwaige Betriebs- und Geschäftsgeheimnisse konkret. Bitte bestätigen Sie umgehend, ob bereits ein Zuschlag erteilt wurde. Über eine Beiladung und einen Termin wird gesondert entschieden.

Dr. Marie Keller
Hauptamtliche Beisitzerin
Vergabekammer Niedersachsen''','NO26-VK031')
    docx('32_sachstand_25092026.docx',CITY,'Sachstand am 25 September','25.09.2026, 14:00 Uhr MESZ',[
        ('1 Eingang und Zuschlagsstand','Die Mitteilung der Vergabekammer mit dem fallinternen Geschäftszeichen VK-NO26-L430-N1 ging heute um 09:10:05 Uhr elektronisch ein. Der Abruf im Behördenpostfach erfolgte um 09:13:22 Uhr. Die Stadt hat keinen Zuschlag erteilt. Der für heute frühestens vorgesehene Vertragsschluss wurde nicht ausgelöst. Die Zuschlagsfunktion ist intern gesperrt; die Sachbearbeitung wurde informiert.'),
        ('2 Bearbeitungsstand','Die Vergabekammer hat bis 30.09.2026, 12:00 Uhr MESZ, die Akte und eine Stellungnahme angefordert. Die Antragsschrift und die fünfzehn Anlagen sind in der Akte abgelegt. Der Schriftsatz der Stadt ist noch nicht gefertigt. Es gibt noch keinen Beschluss, keine Beiladungsentscheidung und keinen Verhandlungstermin. Die Bindefrist der Angebote läuft unverändert bis 30.10.2026.'),
        ('3 Zusammenstellung','Gesichert sind die ursprünglichen Angebote, die Bekanntmachung, die Unterlagenfassung, die Antwort an alle Unternehmen, die spätere Aufklärung, die Wertungsunterlagen und die Versanddaten. Die Generalunternehmerbestätigung und die Herstellerunterlage liegen jeweils als eigenständige Dokumente vor. Ein Inbetriebnahmeprotokoll der Referenzanlage sowie eine vor Angebotsabgabe schriftlich festgelegte Ventilatorauswahl liegen weiterhin nicht vor. Die unterschiedlichen Einschätzungen im Wertungsvermerk bleiben unverändert dokumentiert.')], 'Mara Seidel, Zentrale Vergabestelle.')
    print(json.dumps({'case':str(CASE),'netto':[a,b,c],'originale_bisher':len([p for p in CASE.iterdir() if p.name[:2].isdigit()])},ensure_ascii=False))

if __name__ == '__main__':
    main()
