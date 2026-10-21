#!/usr/bin/env python3
"""Prüft Bestandsschutz, native Ergänzungen und konkrete Buchungsabgrenzungen."""
import hashlib
import json
import re
import unittest
from collections import Counter
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
FIX=ROOT/'scripts/fixtures/fintech-bremen'
CASE=ROOT/'testakten/fintech-darlehen-vertragsuebernahme-bremen'
DATA=json.loads((FIX/'anfechtung-ergaenzung.json').read_text())
OUT=CASE/DATA['folder']


class Anfechtungsvertiefung(unittest.TestCase):
    def test_all_42_originals_preserved(self):
        expected=json.loads((FIX/'anfechtung-original42-sha256.json').read_text())
        self.assertEqual(len(expected),42)
        for name,digest in expected.items():
            with self.subTest(file=name): self.assertEqual(hashlib.sha256((CASE/name).read_bytes()).hexdigest(),digest)

    def test_twenty_files_and_readme_inventory(self):
        expected={d['id']+'.'+d['format'] for d in DATA['documents']}
        expected.update(d['id']+'.xlsx' for d in DATA['workbooks'])
        expected.update(d['id']+'.eml' for d in DATA['emails'])
        self.assertEqual({f.name for f in OUT.iterdir()},expected)
        self.assertEqual(Counter(f.suffix for f in OUT.iterdir()),{'.pdf':4,'.docx':4,'.xlsx':2,'.eml':10})
        readme=(CASE/'README.md').read_text()
        self.assertIn('62 eigenständige Arbeitsdateien',readme)
        for name in expected: self.assertEqual(readme.count(f']({DATA["folder"]}/{name})'),1)

    def test_full_prose_and_real_mime_attachments(self):
        count=0
        for d in DATA['emails']:
            m=BytesParser(policy=policy.default).parsebytes((OUT/(d['id']+'.eml')).read_bytes())
            self.assertEqual(m.defects,[])
            self.assertEqual(m['Date'].datetime.isoformat(),d['date'])
            self.assertEqual(str(m['Subject']),d['subject'])
            self.assertGreater(len(m.get_body(preferencelist=('plain',)).get_content().split()),130)
            for header in ['From','To','Cc']:
                for address in re.findall(r'[\w.-]+@([\w.-]+)',str(m.get(header,''))): self.assertTrue(address.endswith('.example'))
            parts=list(m.iter_attachments())
            self.assertEqual([p.get_filename() for p in parts],d['attachments'])
            for part in parts:
                count+=1
                self.assertEqual(part.get_content_disposition(),'attachment')
                self.assertEqual(part.get_payload(decode=True),(OUT/part.get_filename()).read_bytes())
        self.assertEqual(count,11)

    def test_documents_are_filled_and_separate_from_existing_pleadings(self):
        for d in DATA['documents']:
            p=OUT/(d['id']+'.'+d['format'])
            if d['format']=='pdf':
                reader=PdfReader(p)
                self.assertEqual(len(reader.pages),1)
                text='\n'.join(page.extract_text() for page in reader.pages)
            else:
                doc=Document(p)
                text='\n'.join(p.text for p in doc.paragraphs)
                self.assertEqual(doc.styles['Normal'].font.name,'Times New Roman')
                self.assertEqual(doc.styles['Normal'].font.size.pt,11)
                self.assertEqual(doc.core_properties.author,'Klotzkette')
            self.assertGreater(len(text.split()),240)
            self.assertNotRegex(text,r'\bMusterlösung\b|\bTestakte\b|generated with AI|\[TODO\]')
        text=json.dumps(DATA,ensure_ascii=False)
        for fragment in ['keine Ausgleichsbuchung','nicht ausgeführt','keine Wortabschrift','noch nicht versandter','keine abschließende']:
            self.assertIn(fragment,text)

    def test_numeric_dates_and_preserved_formula_values(self):
        expected={
          '59_Stichtage_und_Belegluecken_20260924.xlsx':('Stichtage',{'B7':478380,'B10':42520,'C8':598600,'C9':1076980,'C10':-693980,'B12':'offen','C12':'offen'}),
          '60_Zahlungswege_und_Vorschlaege_20260924.xlsx':('Abgleich',{'B6':590000,'B7':515000,'B8':505000,'B9':500000,'B10':42840,'B11':500000,'B12':19})
        }
        for name,(sheet,values) in expected.items():
            cached=load_workbook(OUT/name,data_only=True)
            formulas=load_workbook(OUT/name,data_only=False)
            for ref,value in values.items():
                self.assertEqual(cached[sheet][ref].value,value)
                self.assertEqual(formulas[sheet][ref].data_type,'f')
            for ws in cached:
                self.assertEqual(ws.page_setup.orientation,'landscape')
                self.assertFalse(any(c.data_type=='e' for row in ws for c in row))
            if sheet=='Stichtage':
                self.assertIsNone(cached[sheet]['B11'].value)
                self.assertIsNone(cached[sheet]['C11'].value)
                self.assertEqual(cached['Belege']['C6'].value.date().isoformat(),'2024-03-12')
            else:
                self.assertEqual(cached['Bewegungen']['B6'].value.date().isoformat(),'2022-11-15')
                self.assertEqual(cached['Bewegungen']['F27'].value,'nicht ausgeführt')


if __name__=='__main__': unittest.main()
