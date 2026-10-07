#!/usr/bin/env python3
"""Erzeugt 24 kurze Fachanwaltsakten mit je neun Originalen; XLSX baut das Partnerskript.

Inhalte und Honorarparameter: si-native-kanzlei-faelle.json. Bestehende Excel-Dateien,
Gesamt-PDFs und Dateien anderer Akten werden nicht verändert. --case SLUG begrenzt.
"""
from pathlib import Path
from datetime import datetime, timezone
from email.message import EmailMessage
from email.headerregistry import Address
from email.policy import SMTP
from email.parser import BytesParser
from email.utils import format_datetime
from xml.sax.saxutils import escape
import argparse, hashlib, importlib.util, json, re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from pypdf import PdfReader
from akten_build_runtime import serif_font_path
ROOT=Path(__file__).resolve().parents[1]
STAMP=datetime(2026,10,7,10,0,tzinfo=timezone.utc)
from testakte_disclaimer import NOTICE_MARKDOWN as NOTICE
NAMES=['01_Mandatsanfrage.eml','02_Rueckfragen_und_Honorar.eml','03_Gegenseite_oder_Fachstelle.eml','04_Zeitnotiz.eml','05_Vertrags_oder_Stammunterlage.pdf','06_Ereignis_oder_Bescheid.pdf','07_Beleg_und_Nachweis.pdf','08_Zeiten_und_Honorar.xlsx','09_Fachlicher_Dokumententwurf.docx','10_Mandats_und_Honorarvermerk.docx']
def euro(n):return f'{n:,.2f}'.replace(',','X').replace('.',',').replace('X','.')+' EUR'
def fee(c):
 m=c['fee_model']
 if m=='hourly':return f"Vorgesehen ist ein Stundenhonorar von {euro(c['rate'])} netto. Tatsächliche Minuten werden ohne automatische Aufrundung erfasst; ein Budgetdeckel ist nicht vereinbart."
 if m=='capped':return f"Vorgesehen sind {euro(c['rate'])} netto je Stunde mit einem verbindlichen Honorarhöchstbetrag von {euro(c['cap'])} netto für den abgegrenzten Auftrag. Der Deckel wird durch weitere Zeiteinträge nicht automatisch erhöht."
 if m=='flat':return f"Für den abgegrenzten Auftrag ist ein Festhonorar von {euro(c['flat_fee'])} netto vorgesehen. Zeiten werden zur internen Steuerung erfasst und nicht zusätzlich als Stundenhonorar berechnet."
 if m=='estimate':return f"Vorgesehen sind {euro(c['rate'])} netto je Stunde. Die bisherige Kostenschätzung beträgt {euro(c['estimate'])} netto und ist kein verbindlicher Deckel. Bei absehbarer Überschreitung wird vor weiterem Aufwand eine Abstimmung vorbereitet."
 return 'Die Vergütung richtet sich nach dem RVG. Gebührengegenstand, Angelegenheit, Verfahrensstand und etwaige Anrechnung müssen vor der Abrechnung geklärt werden. Erfasste Minuten sind kein zusätzliches Stundenhonorar.'
def agreement(c):
 return 'Die Mandantschaft hat diese Honorargrundlage im Gespräch am 6. Oktober bestätigt; der vorliegende Vermerk hält den Gesprächsstand fest. Form, Pflichtinformationen und Reichweite bleiben vor Rechnungsfreigabe anhand der vollständigen Mandatsunterlagen zu prüfen.' if c['agreed'] else 'Diese Honorargrundlage ist bislang nur vorgeschlagen und von der Mandantschaft nicht bestätigt. Die Bearbeitung erfolgt vorläufig zur Sachverhalts und Fristklärung; eine Rechnung aus dem vorgeschlagenen Modell ist nicht freigegeben.'
def mail(path,c,n,subject,body,sender,recipient):
 msg=EmailMessage(policy=SMTP)
 display,address=sender.rsplit(' <',1);msg['From']=Address(display_name=display,addr_spec=address.rstrip('>'))
 display,address=recipient.rsplit(' <',1);msg['To']=Address(display_name=display,addr_spec=address.rstrip('>'))
 msg['Subject']=f"{c['case_id']} | {subject}"
 msg['Date']=format_datetime(datetime(2026,10,6 if n<4 else 7,8+n*2,12,tzinfo=timezone.utc))
 msg['Message-ID']=f"<{c['case_id'].lower()}-{n:02}@blum-rapp.example>"
 msg['Content-Language']='de-DE'
 if n==2:msg['In-Reply-To']=f"<{c['case_id'].lower()}-01@blum-rapp.example>"
 msg.set_content(body,charset='utf-8')
 path.write_bytes(msg.as_bytes())
 read=BytesParser(policy=SMTP).parsebytes(path.read_bytes())
 assert read.get_body(preferencelist=('plain',)).get_content().replace('\r\n','\n').strip()==body.strip()
 assert not read.defects
 return {'characters':len(body),'message_id':str(msg['Message-ID'])}
def docx(path,title,text,c):
 doc=Document();sec=doc.sections[0]
 sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
 for sn in ('Normal','Title','Heading 1','Heading 2'):
  st=doc.styles[sn];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
  for script in ('ascii','hAnsi','eastAsia','cs'):st.element.get_or_add_rPr().rFonts.set(qn('w:'+script),'Times New Roman')
  for key in list(st.element.get_or_add_rPr().rFonts.attrib):
   if key.endswith('Theme'):del st.element.get_or_add_rPr().rFonts.attrib[key]
  st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.03
 doc.styles['Title'].font.size=Pt(14);doc.styles['Title'].font.bold=True
 doc.styles['Heading 1'].font.bold=True
 for b in doc.styles.element.xpath('.//w:pBdr'):b.getparent().remove(b)
 doc.add_paragraph('Kanzlei Blum und Rapp | '+c['case_id'])
 doc.add_paragraph('7. Oktober 2026')
 doc.add_paragraph(title,'Title')
 for block in text.split('\n\n'):
  heading=bool(re.match(r'^\d+(?:\.\d+)* ',block)) and '\n' not in block and len(block)<100
  p=doc.add_paragraph(block,'Heading 1' if heading else 'Normal');p.paragraph_format.widow_control=True
  if heading:p.paragraph_format.space_before=Pt(7)
 doc.core_properties.author=doc.core_properties.last_modified_by='Klotzkette'
 doc.core_properties.title=title;doc.core_properties.language='de-DE'
 doc.core_properties.created=doc.core_properties.modified=STAMP
 doc.save(path)
 return {'characters':len(text)}
def pdf(path,item,c):
 body=ParagraphStyle('body',fontName='TNR',fontSize=11,leading=14.1,spaceAfter=8)
 title=ParagraphStyle('title',parent=body,fontName='TNRB',fontSize=14,leading=17,spaceAfter=12)
 flow=[Paragraph(escape(item['author']),body),Paragraph(escape(item['date']),body),Spacer(1,7),Paragraph(escape(item['title']),title)]
 flow += [Paragraph(escape(p),body) for p in item['text'].split('\n\n')]
 def footer(can,doc):
  can.setFont('TNR',9);can.drawString(54,28,c['case_id']);can.drawRightString(A4[0]-54,28,str(doc.page))
 def invariant(*args,**kwargs):kwargs['invariant']=1;return Canvas(*args,**kwargs)
 SimpleDocTemplate(str(path),pagesize=A4,leftMargin=54,rightMargin=54,topMargin=48,bottomMargin=48,title=item['title'],author='Klotzkette').build(flow,onFirstPage=footer,onLaterPages=footer,canvasmaker=invariant)
 assert len(PdfReader(path).pages)==1
 return {'characters':len(item['text']),'pages':1}
def main():
 parser=argparse.ArgumentParser();parser.add_argument('--case',action='append');parser.add_argument('--write-readmes',action='store_true',help='Bestehende README mit zentralem Downloadblock neu erzeugen');args=parser.parse_args()
 cs=json.loads((ROOT/'scripts/si-native-kanzlei-faelle.json').read_text())
 if args.case:cs=[c for c in cs if c['slug'] in args.case]
 pdfmetrics.registerFont(TTFont('TNR',str(serif_font_path())))
 pdfmetrics.registerFont(TTFont('TNRB',str(serif_font_path(bold=True))))
 manifest=[]
 for c in cs:
  folder=ROOT/'testakten'/c['slug'];folder.mkdir(parents=True,exist_ok=True)
  client=f"{c['client']} <mandat@{c['case_id'].lower()}.example>"
  opp=f"{c['opponent']} <kontakt@gegenseite-{c['case_id'].lower()}.example>"
  lawyer='Dr. Frieda Blum <frieda.blum@blum-rapp.example>'
  assistant='Jona Rapp <jona.rapp@blum-rapp.example>'
  bodies=[
   ('Bitte übernehmen Sie die erste Bearbeitung',f"Sehr geehrte Frau Dr. Blum,\n\n{c['opening']}\n\nDie drei Unterlagen liegen im Aktenordner. Bitte sagen Sie mir, welche Angaben noch fehlen und welches Dokument als Nächstes sinnvoll ist. Ich möchte einen Entwurf sehen, bevor etwas nach außen geht. Über weitere Schritte und die Kosten möchte ich rechtzeitig informiert werden.\n\nMit freundlichen Grüßen\n{c['client']}",client,lawyer),
   ('Rückfragen und Honorar für den abgegrenzten Auftrag',f"Guten Tag,\n\n{c['questions']}\n\n{fee(c)}\n\n{agreement(c)} Die genannten Beträge sind netto; Umsatzsteuer und mögliche Auslagen sind vor der verbindlichen Verbraucherinformation und der Rechnung gesondert auszuweisen. Bitte bestätigen Sie zusätzliche Aufgaben ausdrücklich. Den ersten Dokumententwurf erhalten Sie zur inhaltlichen Abstimmung.\n\nMit freundlichen Grüßen\nDr. Frieda Blum",lawyer,client),
   ('Aktueller Stand und noch offene Unterlagen',f"Guten Tag,\n\n{c['response']}\n\nBitte verwenden Sie bei Ihrer Antwort das bekannte Vorgangszeichen und benennen Sie die konkreten Unterlagen, auf die Sie sich beziehen. Wir haben mit dieser Nachricht keine weitergehende Einigung erklärt.\n\nMit freundlichen Grüßen\n{c['opponent']}",opp,lawyer),
  ]
  t=c['times'];body=f"Hallo Jona,\n\nbitte ordne diese Tätigkeiten ausschließlich der Akte {c['case_id']} zu. Die ersten beiden Zeitangaben sind von mir bestätigt:\n\n1 Am {t[0]['date']} habe ich {t[0]['minutes']} Minuten für folgendes Narrative aufgewendet: {t[0]['narrative']}.\n\n2 Am {t[1]['date']} habe ich {t[1]['minutes']} Minuten für folgendes Narrative aufgewendet: {t[1]['narrative']}.\n\n3 Für das Telefonat am {t[2]['date']} fehlen Dauer und endgültiger Rechnungstext. Bitte gezielt nachfragen und bis dahin keinen geschätzten Betrag buchen.\n\n4 Deine Ablagenotiz nennt {t[3]['minutes']} Minuten. Ob diese organisatorische Tätigkeit nach der vereinbarten Grundlage abrechenbar ist, ist noch offen. Sie bleibt unbestätigt und vorläufig nicht abrechenbar.\n\nBitte den internen Rechnungsentwurf nach der Honorargrundlage in Dokument 10 fortschreiben. Es gibt noch keine freigegebene Rechnungsnummer und keinen Versand. Festpreis, Deckel und RVG dürfen nicht durch bloßes Multiplizieren der gesamten Zeit ersetzt werden.\n\nDanke\nFrieda"
  bodies.append(('Zeiten eintragen und offene Dauer rückfragen',body,lawyer,assistant))
  for idx,(subject,body,sender,recipient) in enumerate(bodies,1):
   name=NAMES[idx-1];meta=mail(folder/name,c,idx,subject,body,sender,recipient);manifest.append(dict(case=c['slug'],name=name,kind='eml',title=subject,**meta))
  for idx,item in enumerate(c['evidence'],5):
   name=NAMES[idx-1];meta=pdf(folder/name,item,c);manifest.append(dict(case=c['slug'],name=name,kind='pdf',title=item['title'],**meta))
  name=NAMES[8];meta=docx(folder/name,(c['draft_type'] if c['draft_type'].lower().endswith('entwurf') else c['draft_type']+' Entwurf'),c['draft'],c);manifest.append(dict(case=c['slug'],name=name,kind='docx',title=c['draft_type'],**meta))
  honorar=f"1 Mandat und Auftrag\n\nMandantin beziehungsweise Mandant ist {c['client']}. Gegenüber steht {c['opponent']}. Gegenstand ist: {c['issue']}. Der erste Arbeitsschritt ist die Vorbereitung des folgenden Dokuments: {c['draft_type']}. Gerichtliche Erweiterungen, Anerkenntnisse, Vergleiche und externer Versand sind daraus nicht automatisch freigegeben.\n\n2 Honorarstand\n\n{fee(c)} {agreement(c)} Die genannten Geldbeträge verstehen sich netto. Umsatzsteuer, Auslagen und erforderliche Verbraucherinformationen sind gesondert zu prüfen.\n\n3 Leistung und Rechnung\n\nDie bestätigten Tätigkeiten und offenen Rückfragen stehen in der Zeitnotiz vom 7. Oktober und in der Excel-Arbeitsmappe. Fehlende Minuten sind nicht null. Unbestätigte Zeiten bleiben von einem abrechenbaren Zwischensaldo getrennt. Ein Entwurf erhält noch keine endgültige Rechnungsnummer. Zahlungsstatus, E-Rechnungsformat, Rechnungsadresse und Fälligkeit sind vor Ausgabe zu ergänzen.\n\n4 Mandatsprüfung und nächster Schritt\n\nIdentität, Vertretungsbefugnis, Interessenkonflikte und die fallbezogene Anwendbarkeit des Geldwäschegesetzes sind bei der Aufnahme zu dokumentieren. Der Name des Fachgebiets ersetzt diese Prüfung nicht. Bei Verbrauchern sind Abschlussweg und gegebenenfalls Widerrufsinformationen zu klären. Verantwortlich für die fachliche Freigabe ist Dr. Frieda Blum. Jona Rapp pflegt die Ablage. Nach Ergänzung der offenen Tatsachen wird der Dokumententwurf aus Datei 09 erneut geprüft."
  name=NAMES[9];meta=docx(folder/name,'Vermerk zu Mandat und Honorar',honorar,c);manifest.append(dict(case=c['slug'],name=name,kind='docx',title='Vermerk zu Mandat und Honorar',**meta))
  table='\n'.join(f'| `{n}` | '+([ 'Mandatsanfrage','Fachliche Rückfragen und Honorarstand','Position der Gegenseite oder Fachstelle','Bestätigte Zeiten und zwei offene Einträge',c['evidence'][0]['title'],c['evidence'][1]['title'],c['evidence'][2]['title'],'Erklärte Zeiten und Honorarberechnung',c['draft_type']+' mit auszufüllenden Stellen','Interner Mandats und Honorarstand'][i])+' |' for i,n in enumerate(NAMES))
  base='https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/si-native-kanzlei-v445.33.3/'
  readme=f"# {c['title']}\n\nKleine Testakte für **{c['fachgebiet']}** im Plugin [KI-native Kanzlei](../../ki-native-kanzlei/README.md). Aktenstand: 7. Oktober 2026. Aktenzeichen: `{c['case_id']}`.\n\n## 1 Einstieg\n\n{c['issue']}.\n\n{c['task']} Der fachliche Entwurf in Datei 09 ist absichtlich unvollständig: Er enthält ausdrücklich bestellte, beschriftete Eingabestellen und darf erst nach fachlicher Prüfung finalisiert werden. Er ist keine Musterlösung.\n\n## 2 Zehn Originaldateien\n\n| Datei | Inhalt |\n|---|---|\n{table}\n\nDie Akte enthält genau vier E Mails, drei PDF Belege, eine Excel-Arbeitsmappe und zwei Word Dateien. Die Excel-Arbeitsmappe trennt bestätigte Werte, fehlende Angaben und einen noch nicht freigegebenen Rechnungsentwurf. Die vier Zeitzeilen stimmen mit Datei 04 und der Honorargrundlage in Datei 10 überein.\n\n## 3 Downloads\n\n{NOTICE}\n\n| Fassung | Download |\n|---|---|\n| Originaldateien | [Akten-ZIP]({base}testakte-{c['slug']}.zip) |\n| Je Dokument ein PDF | [Einzel-PDF-ZIP]({base}testakte-{c['slug']}-einzelpdfs.zip) |\n| Lesefassung | [Gesamt-PDF](gesamt-pdf/{c['slug']}_gesamt.pdf) |\n\n## 4 Herkunft und Grenzen\n\nAlle Personen, Unternehmen, Dokumente, Kontaktdaten und Sachverhalte sind erfunden. Reale Städtenamen dienen nur der Einordnung. Genannte Behörden sind keine Urheber tatsächlich ergangener Entscheidungen. Reservierte `.example` Adressen verhindern eine versehentliche Kontaktaufnahme.\n\n<!-- reserved-example-contacts -->\n\nDie Akte ist ein kurzer Einstieg in Mandatsaufnahme, fachliche Dokumentbearbeitung, Zeiterfassung und Honorarsteuerung. Fehlende Originale werden ausdrücklich als fehlend behandelt. Es werden keine tatsächlichen Konten angelegt, Nachrichten versandt oder behördlichen Vorgänge eröffnet. Die Rechtsquellen stehen im [Plugin](../../ki-native-kanzlei/references/rechtsquellen.md); zusätzliche fallbezogene Normen sind in den Entwürfen genannt und vor praktischer Verwendung aktuell zu prüfen.\n"
  # Gepflegte Downloadseiten standardmäßig erhalten; beim Neubau zentrale Marker nutzen.
  readme_path=folder/'README.md'
  if not readme_path.exists() or args.write_readmes:
   spec=importlib.util.spec_from_file_location('si_case_downloads',ROOT/'scripts/inject-gesamt-pdf-section.py')
   mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
   start=readme.index('## 3 Downloads');end=readme.index('## 4 Herkunft')
   readme=readme[:start]+readme[end:].replace('## 4 Herkunft','## 3 Herkunft',1)
   heading,rest=readme.split('\n',1)
   block=mod.section_block(c['slug'],f"gesamt-pdf/{c['slug']}_gesamt.pdf",has_einzelpdf=True)
   block=block.replace('Die ZIP-Links laden den zuletzt veröffentlichten Release.','Die ZIP-Links laden den in der Release-Konfiguration festgelegten Aktenrelease.')
   block=block.replace('ZIP links refer to the latest published release.','ZIP links refer to the release selected in the repository asset configuration.')
   readme_path.write_text(heading+'\n\n'+block+'\n'+rest.lstrip())
  rubric={'name':c['title'],'plugin':'si-native-kanzlei','stand':'2026-10-07','checks':[{'id':'arbeitsbestand','check_type':'working_file_count','description':'Zehn native Originale sind vorhanden.','min':10}]+[{'id':f'original-{i:02}','check_type':'file_exists','description':f'Das Aktenstück {i:02} ist vorhanden.','path':name} for i,name in enumerate(NAMES,1)]+[{'id':'fachlicher-entwurf','check_type':'human_review','description':'Der fachliche Entwurf bearbeitet den konkreten Sachverhalt, enthält benannte Blanks und bleibt vor Freigabe ein Entwurf.'},{'id':'honorar-und-zeit','check_type':'human_review','description':'Honorargrundlage, bestätigte Zeiten und ungeklärte Einträge werden unterschieden; ein unbestätigtes Angebot wird nicht als Rechnung versandt.'},{'id':'quellen-und-anlagen','check_type':'human_review','description':'Behauptungen, Belege und fehlende Originale bleiben unterscheidbar; Anlagen werden passend zum fertigen Entwurf ausgewählt.'}]}
  # JSON ist gültiges YAML; der Rubric-Loader verwendet YAML safe_load.
  (folder/'rubric.yaml').write_text(json.dumps(rubric,ensure_ascii=False,indent=2)+'\n')
 for m in manifest:m['sha256']=hashlib.sha256((ROOT/'testakten'/m['case']/m['name']).read_bytes()).hexdigest()
 qa=ROOT/'quality/si-native-kanzlei';qa.mkdir(parents=True,exist_ok=True)
 old=[]
 if args.case and (qa/'testakten-originale.json').exists():old=[x for x in json.loads((qa/'testakten-originale.json').read_text()) if x['case'] not in args.case]
 (qa/'testakten-originale.json').write_text(json.dumps(sorted(old+manifest,key=lambda x:(x['case'],x['name'])),ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({'cases':len(cs),'originals':len(manifest),'eml':sum(x['kind']=='eml' for x in manifest),'pdf':sum(x['kind']=='pdf' for x in manifest),'docx':sum(x['kind']=='docx' for x in manifest),'xlsx':'separates Partnerskript'},ensure_ascii=False))
if __name__=='__main__':main()
