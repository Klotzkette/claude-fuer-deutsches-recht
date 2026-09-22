#!/usr/bin/env python3
"""Erzeugt nur die 27 nativen Lindenhof-Quellen; keine Release-Exporte.

Aufruf aus beliebigem Verzeichnis: python3 scripts/build-strassennutzung-lindenhof-akte.py
Abhängigkeiten: python-docx, reportlab, pypdf. Keine Tabellenkalkulation erforderlich.
Der fachliche Kanon und sämtliche Texte stehen bewusst konkret in dieser Datei.
"""

from __future__ import annotations

import argparse
import csv
import io
from datetime import datetime
from decimal import Decimal
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph
from readme_decimal_headings import normalize_decimal_headings


ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "testakten/strassennutzung-poller-lieferzufahrt-lindenhof-muenster"
AZ = "MS-LB-2026-0417"
FIXED_TIME = datetime(2026, 9, 22, 16, 0)
INVENTORY: dict[str, str] = {}


def contact_text(value: str) -> str:
    domains = {
        "lindenhof-vorrat.example.org": "lindenhof-vorrat-ms.de",
        "stadt-muenster.example.org": "strassenstelle-lindenbogen-ms.de",
        "haus-bruening.example.org": "haus-bruening-ms.de",
        "feuerwehr-muenster.example.org": "einsatzstelle-lindenbogen-ms.de",
        "post.example.org": "post-lindenbogen-ms.de",
        "klenke-metall.example.org": "klenke-metall-ms.de",
        "frischefahrt-west.example.org": "frischefahrt-west-ms.de",
        "quartierbus-muenster.example.org": "quartierbus-west-ms.de",
        "cafe-kranich.example.org": "cafe-kranich-ms.de",
    }
    for original, replacement in domains.items():
        value = value.replace(original, replacement)
    return value

SHOP = ("Lindenhof Vorrat", "Mara Hölscher | Lindenbogen 18 | 48147 Münster", "mara.hoelscher@lindenhof-vorrat.example.org")
CITY = ("Stadt Münster", "Straßenverkehr | Johanna Rensing | Verwaltungsstelle Lindenplatz 4 | 48147 Münster", "johanna.rensing@stadt-muenster.example.org")
WORKS = ("Stadt Münster", "Straßenbetrieb | Tobias Korte | Werkhof Hegenkamp 7 | 48147 Münster", "tobias.korte@stadt-muenster.example.org")
OWNER = ("Elisabeth Brüning", "Grundstück Lindenbogen 18 | vertreten durch Nils Brüning | Weidenstieg 6 | 48147 Münster", "nils.bruening@haus-bruening.example.org")


def remember(name: str, description: str) -> Path:
    INVENTORY[name] = description
    return CASE / name


def stable_docx(document: Document, destination: Path) -> None:
    props = document.core_properties
    props.created = props.modified = FIXED_TIME
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.revision = 1
    buf = io.BytesIO()
    document.save(buf)
    with ZipFile(buf) as original, ZipFile(destination, "w", ZIP_DEFLATED) as target:
        for member in sorted(original.namelist()):
            entry = ZipInfo(member, (2026, 9, 22, 16, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            target.writestr(entry, original.read(member))


def docx_letter(name, sender, recipient, date, subject, reference, paragraphs, closing, attachments="Keine."):
    sender = tuple(contact_text(line) for line in sender)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(19), Mm(19)
    sec.left_margin, sec.right_margin = Mm(24), Mm(22)
    sec.header_distance, sec.footer_distance = Mm(9), Mm(10)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in list(fonts.attrib):
            if "theme" in attr.lower():
                del fonts.attrib[attr]
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attr), "Times New Roman")
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.06
    doc.styles["Title"].font.size = Pt(14)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    doc.styles["Title"].paragraph_format.keep_with_next = True
    head = doc.add_paragraph()
    head.paragraph_format.space_after = Pt(3)
    run = head.add_run(sender[0])
    run.bold, run.font.size = True, Pt(16)
    for line in sender[1:]:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)
        p.runs[0].font.size = Pt(9)
    p = doc.add_paragraph(recipient)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph(f"Münster, {date}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph(subject, style="Title")
    p = doc.add_paragraph(reference)
    p.paragraph_format.space_after = Pt(10)
    p.runs[0].font.size = Pt(9)
    for paragraph in paragraphs:
        p = doc.add_paragraph(paragraph)
        p.paragraph_format.widow_control = True
    p = doc.add_paragraph(closing)
    p.paragraph_format.keep_together = True
    p = doc.add_paragraph("Anlagen: " + attachments)
    p.runs[0].font.size = Pt(9)
    p.paragraph_format.keep_together = True
    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    footer.add_run(f"{reference.split(' | ')[0]}  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    doc.core_properties.title = subject
    stable_docx(doc, remember(name, subject))


def build_letters():
    docx_letter("01_antrag_hoelscher_2026-07-16.docx", SHOP,
        "Stadt Münster\nStraßenverkehr\nFrau Johanna Rensing\nVerwaltungsstelle Lindenplatz 4\n48147 Münster", "16.07.2026",
        "Zufahrt zum Hof am Lindenbogen 18", "LHV-16/07 | Ihr Gesprächsangebot vom 13.07.2026", [
        "Sehr geehrte Frau Rensing,",
        "ich beantrage, dass die angekündigte Pollersperre die Zufahrt zu unserem Hof für Lieferungen weiterhin ermöglicht. Ich benötige dafür einen Schlüssel und eine Öffnungsmöglichkeit montags bis freitags von 06:00 bis 18:00 Uhr. Der Verkauf beginnt um 08:00 Uhr; Frischware und Getränkekisten werden auch nachmittags nachgeliefert. Ich beantrage keine Abstellfläche auf der Straße und keinen Parkplatz für unsere Kunden.",
        "Seit Mai 2024 betreibe ich im Erdgeschoss einen Lebensmittel- und Haushaltswarenladen mit 58 m² Verkaufsfläche und 19 m² Lager. Mein Mitarbeiter Deniz Arslan nimmt morgens Ware an, während ich im Verkauf bin. Die Kühlware kommt überwiegend in einem 3,5-t-Transporter. Getränke Röttger nutzt teils einen 7,5-t-Lkw; dieser entlädt bislang vor dem Tor, weil er wegen seiner Höhe nicht in den Hof passt. Die übrigen Transporter fahren zum Entladen hinein.",
        "Die Ladentür ist 0,96 m breit. Zwischen Verkaufsraum und Lager liegen drei Stufen. Kisten vom Lindenplatz durch den Verkauf zu tragen ist deshalb nicht dasselbe wie die Anlieferung über das ebene Hoftor. Über den Eschenring sind wir mit unserem eigenen kurzen Transporter schon gefahren. Für den Getränkelastwagen halte ich die enge Einmündung nicht für befahrbar; eine Vermessung habe ich nicht.",
        "Unsere Vermieterin hat uns den Hof zur Warenannahme mitvermietet. Ich berufe mich insoweit auch auf den Straßenanliegergebrauch nach Paragraf 14a des Straßen- und Wegegesetzes NRW. Ob Ihre Planung einen anderen Zufahrtsweg vorsieht, konnte ich aus dem Gespräch nicht erkennen. Bitte beziehen Sie mich vor dem Einbau ein und teilen Sie mir mit, wie Lieferanten außerhalb einer festen Morgenstunde an das Grundstück gelangen sollen.",
        "Im Hof stehen keine Kundenparkplätze zur Verfügung. Auch ich ärgere mich über Autos vor dem Tor. Ich kann aber nicht bei jedem abgestellten Fahrzeug erkennen, welchen Laden die Fahrerin oder der Fahrer besucht. Eine neue Zufahrt oder eine Änderung des Bordsteins plane ich nicht."
    ], "Mit freundlichen Grüßen\nMara Hölscher\nInhaberin", "Nachtrag zur Hofnutzung vom 18.04.2024, eine Ausfertigung.")

    docx_letter("02_mietnachtrag_hof_2024-04-18.docx", OWNER,
        "Frau Mara Hölscher\nLindenhof Vorrat\nLindenbogen 18\n48147 Münster", "18.04.2024",
        "Nachtrag zur Nutzung des rückwärtigen Hofs", "Mietvertrag LB18-2024 | Mietbeginn 01.05.2024", [
        "Sehr geehrte Frau Hölscher,",
        "wir halten die am 17. April besprochene Ergänzung unseres Mietvertrags vom 4. April 2024 wie folgt fest. Sie dürfen den südlich an die Durchfahrt anschließenden Hof von 7,20 m Breite und 9,00 m Tiefe während der Mietzeit zur Warenannahme für Ihren Laden mitbenutzen. Der Hof wird gemeinsam mit den Bewohnern der beiden Obergeschosse genutzt. Eine ausschließliche Stellplatzvermietung ist damit nicht verbunden.",
        "Transporter dürfen zum unmittelbaren Entladen durch das 2,85 m breite Tor einfahren, sofern sie die lichte Höhe von 3,10 m einhalten. Die Torflügel öffnen nach innen. Nach dem Entladen sind sie wieder zu schließen. Der 1,20 m breite Weg entlang der östlichen Hauswand zur hinteren Wohnungstreppe bleibt frei. Lieferfahrzeuge dürfen dort nicht über Nacht stehen. Sie geben diesen Hinweis auch an wechselnde Fahrer weiter.",
        "Die Mitbenutzung ist in der vereinbarten monatlichen Netto-Kaltmiete von 1.180,00 EUR enthalten. Die Betriebskostenvorauszahlung beträgt unverändert 220,00 EUR monatlich. Ein zusätzliches Entgelt für die Hofnutzung wird nicht erhoben. Die übrigen Regelungen des Mietvertrags bleiben unverändert bestehen.",
        "Im Hof und in der Durchfahrt dürfen keine Kundenfahrzeuge abgestellt werden. Eine öffentliche Parkfläche vor dem Grundstück ist nicht Gegenstand des Mietvertrags. Frau Brüning übernimmt keine Zusage dazu, ob oder in welchem Umfang die Stadt die Zufahrt von der öffentlichen Straße künftig beschränkt. Bauliche Veränderungen am Tor, an der Entwässerung oder am Gehweg bedürfen einer gesonderten Abstimmung.",
        "Am 18. April haben Sie zwei Hoftorschlüssel übernommen. Es handelt sich um Schlüssel des privaten Tors, nicht um Schlüssel für Einrichtungen im Straßenraum. Die Annahme dieses Nachtrags haben Sie heute um 14:15 Uhr per E-Mail bestätigt. Wir legen diese Korrespondenz zur Mietvertragsakte; die vorliegende Ausfertigung gibt den vereinbarten Text wieder."
    ], "Mit freundlichen Grüßen\nNils Brüning\nfür Elisabeth Brüning", "Keine.")

    docx_letter("03_mitteilung_fahrradstrasse_2026-07-24.docx", CITY,
        "Frau Mara Hölscher\nLindenhof Vorrat\nLindenbogen 18\n48147 Münster", "24.07.2026",
        "Verkehrsumstellung im Lindenbogen", f"{AZ} | Ihr Antrag vom 16.07.2026", [
        "Sehr geehrte Frau Hölscher,",
        "die Einrichtung der Fahrradstraße im Lindenbogen zwischen Lindenplatz und Eschenring ist für den 24. August 2026 vorgesehen. An beiden Zufahrten werden Zeichen 244.1 mit dem Zusatzzeichen „Anlieger frei“ aufgestellt. Bei Station 64 m ist eine Polleranlage mit herausnehmbarem Mittelpfosten vorgesehen. Die östliche Zufahrt über den Eschenring bleibt baulich offen. Die Anlage gehört zur verkehrsrechtlichen Anordnung SV-LB-26-118 vom 23. Juli 2026 nach Paragraf 45 StVO.",
        "Der Mittelpfosten soll die durchgehende Abkürzung zwischen Lindenplatz und Eschenring unterbrechen. Er ist nicht unmittelbar vor Ihrem Tor vorgesehen, sondern 32 m westlich davon. Nach Ausbau des Mittelpfostens beträgt die Durchfahrtsbreite zwischen den festen Seitenpfosten nach Planung 3,20 m. Wir gehen bisher davon aus, dass die Belieferung kleinerer Gewerbebetriebe von Osten und ergänzend über einen geregelten Schließdienst möglich ist.",
        "Ihre Angaben zu dem Getränkelastwagen und zu den Stufen im Laden nehmen wir zum Anlass, einen befristeten Schlüsselversuch vorzubereiten. Eine ganztägig offene Durchfahrt ist nicht vorgesehen. Bitte benennen Sie die regelmäßig eingesetzten Fahrzeuge, die zuständige Person vor Ort und die Zeiten, zu denen diese Person den Pfosten selbst wieder einsetzen kann. Ein Schlüssel soll nicht bei einem ständig wechselnden Fahrerkreis verbleiben.",
        "Die an der Nordseite vorhandenen markierten Parkstände bleiben bestehen. Vor dem südlichen Hoftor wird kein zusätzlicher Parkstand eingerichtet. Die Haltestelle Eschenring liegt außerhalb des Lindenbogens; der Linienbus benutzt den Lindenbogen nicht. Für Einsatzfahrzeuge wird der Straßenbetrieb eine gesonderte Öffnungsmöglichkeit abstimmen.",
        "Mit diesem Schreiben ist über Ihren Antrag auf die beantragten Öffnungszeiten noch nicht abschließend entschieden. Auch die von Ihnen angesprochene Grundstücksnutzung wird damit nicht bewertet. Herr Korte meldet sich wegen der Schlüsselübergabe und der Einweisung, sobald die Anlage montiert ist."
    ], "Mit freundlichen Grüßen\nJohanna Rensing\nStraßenverkehr", "Keine.")

    docx_letter("04_schluesselversuch_2026-08-20.docx", CITY,
        "Frau Mara Hölscher\nLindenhof Vorrat\nLindenbogen 18\n48147 Münster", "20.08.2026",
        "Befristeter Schließdienst für die Lieferzufahrt", f"{AZ} | Ergänzung unserer Mitteilung vom 24.07.2026", [
        "Sehr geehrte Frau Hölscher,",
        "für den Zeitraum vom 24. August bis einschließlich 15. Oktober 2026 stellen wir Ihnen den Schlüssel LB-17 zur Verfügung. Im Rahmen dieses Versuchs dürfen Sie oder Ihr benannter Mitarbeiter Deniz Arslan den Mittelpfosten P1 montags bis freitags zwischen 06:30 und 10:00 Uhr für einzelne Anlieferungen zu Ihrem Betrieb herausnehmen. Außerhalb dieser Zeiten ist im Versuch keine Öffnung durch Ihren Betrieb vorgesehen.",
        "Die Person, die öffnet, bleibt in Sichtweite der Anlage, bis das Lieferfahrzeug passiert hat, und setzt den Pfosten anschließend wieder ein. Müssen Sie für die Ausfahrt erneut öffnen, erfolgt dies gesondert. Eine durchgehend offene Lieferstunde ist nicht vereinbart. Der ausgebaute Pfosten ist in der Halterung am südlichen Rand abzustellen, nicht quer auf dem Gehweg. Der Schlüssel darf weder im Schloss stecken bleiben noch an Lieferanten oder Kunden weitergegeben werden.",
        "Bitte protokollieren Sie verspätete Lieferungen, Wartezeiten und Schwierigkeiten bei der östlichen Einfahrt. Dieser Versuch gestattet kein Halten oder Parken unabhängig von der vorhandenen Beschilderung und hält auch keine Fläche vor dem Hoftor frei. Für die Öffnung bei Einsätzen bestehen gesonderte Zugänge, die nicht von Ihrer Anwesenheit im Laden abhängen sollen.",
        "Der Straßenbetrieb gibt LB-17 am 21. August um 08:20 Uhr im Werkhof aus und erklärt das Einsetzen sowie die Verriegelung. Die Funktionskontrolle erfolgt ohne Lieferfahrzeug. Bitte melden Sie Schwergängigkeit sofort an Herrn Korte. Einen Ersatzschlüssel für die tägliche Betriebsorganisation sieht der Versuch bislang nicht vor.",
        "Über die von Ihnen beantragte Zeitspanne 06:00 bis 18:00 Uhr entscheiden wir nach Auswertung der Erfahrungen. Erfolgt keine Anschlussregelung, geben Sie LB-17 bis zum 16. Oktober 2026 um 12:00 Uhr am Werkhof zurück. Eine Verlängerung entsteht nicht allein dadurch, dass die Auswertung noch läuft. Für Ihre Rückmeldung wird ein gesonderter Termin angeboten."
    ], "Mit freundlichen Grüßen\nJohanna Rensing\nStraßenverkehr", "Keine.")

    docx_letter("05_anhoerung_2026-09-14.docx", CITY,
        "Frau Mara Hölscher\nLindenhof Vorrat\nLindenbogen 18\n48147 Münster", "14.09.2026",
        "Anhörung zum weiteren Betrieb der Polleranlage", f"{AZ} | Schlüssel LB-17 und Ihr Zusatzantrag vom 27.08.2026", [
        "Sehr geehrte Frau Hölscher,",
        "wir erwägen, den Schlüsselversuch am 15. Oktober 2026 auslaufen zu lassen und weder den beantragten zweiten Schlüssel noch die verlängerten Öffnungszeiten einzuräumen. Vor einer Entscheidung erhalten Sie Gelegenheit, bis zum 30. September 2026 zu den folgenden Beobachtungen Stellung zu nehmen und Ihre Lieferunterlagen zu ergänzen. Dieses Schreiben enthält noch keine abschließende Ablehnung.",
        "Am 4. September blieb der Mittelpfosten nach den bisher vorliegenden Meldungen über die einzelne Lieferung hinaus entfernt. Frau Wegner hat zwischen 10:12 und 10:18 Uhr zwei Pkw durchfahren sehen. Ob diese zum Laden oder zu anderen Grundstücken fuhren, ist nicht geklärt. Für uns stellt sich dennoch die Frage, ob die vereinbarte persönliche Betreuung der Öffnung im laufenden Verkauf dauerhaft funktioniert. Auch ein weiterer Schlüssel könnte nach unserer bisherigen Einschätzung die Zuordnung der Verantwortung erschweren.",
        "Beim Ortstermin am 9. September ließ sich die östliche Einfahrt mit einem 5,99 m langen Transporter erreichen. Hierauf stützt sich unsere bisherige Annahme, dass der Betrieb nicht ausschließlich auf die Pollerdurchfahrt angewiesen ist. Das längere Getränkefahrzeug wurde dabei nicht erprobt. Die Warte- und Zusatzkosten wollen wir anhand einzelner Fahrten nachvollziehen; eine bloße Gesamtsumme reicht dafür nicht aus.",
        "Die Feuerwehr hat die Öffnung mit dem Einsatzschlüssel geprüft, aber keine Befahrung unter Einsatzbedingungen vorgenommen. Der Busbetrieb weist auf Rückstau und Lieferhalte am Eschenring hin. Frau Aydin meldet Schwierigkeiten durch quer abgestellte Pfosten und abgestellte Kisten. Diese Belange sprechen aus unserer Sicht gegen eine Freigabe ohne verlässlichen Schließdienst. Ob eine andere zeitliche oder organisatorische Regelung tragfähig ist, bleibt Gegenstand der Auswertung.",
        "Bitte teilen Sie insbesondere mit, wer zwischen 10:00 und 18:00 Uhr vor Ort öffnen und schließen könnte. Die Besprechung ist am 6. Oktober 2026 um 09:00 Uhr im Besprechungsraum 2, Verwaltungsstelle Lindenplatz 4, vorgesehen. Die schriftliche Frist bleibt davon unberührt. Bis zum 15. Oktober gilt die bisherige Versuchsregelung unverändert; über LB-18 ist noch nicht entschieden."
    ], "Mit freundlichen Grüßen\nJohanna Rensing\nStraßenverkehr", "Ortsterminprotokoll vom 09.09.2026; Stellungnahme Feuerwehr vom 10.09.2026; Busmail vom 08.09.2026.")

    docx_letter("06_ortstermin_2026-09-09.docx", WORKS,
        "An Frau Johanna Rensing, Straßenverkehr\nAbschrift an Mara Hölscher und Nils Brüning", "09.09.2026",
        "Niederschrift zur Zufahrt am Lindenbogen", f"{AZ} | Ortstermin 09.09.2026, 08:30 bis 09:25 Uhr", [
        "Guten Tag Frau Rensing,",
        "am heutigen Termin nahmen Mara Hölscher, Deniz Arslan, Nils Brüning und ich teil. Gemessen wurde mit Bandmaß zwischen sichtbaren Bauteilkanten. Es handelt sich nicht um eine Grundstücksvermessung. Die Fahrbahn ist an der Polleranlage 4,80 m breit. Der nördliche Gehweg misst dort 1,80 m, der südliche 2,00 m. Der herausgenommene Mittelpfosten gibt zwischen den festen Pfosten 3,20 m frei. Eingesetzt verbleiben beidseits je 1,50 m; der Mittelpfosten selbst ist 0,20 m breit.",
        "P1 steht 64 m östlich des Anschlusses Lindenplatz. Das Hoftor liegt weitere 32 m östlich auf der Südseite. Seine lichte Breite beträgt 2,85 m, die lichte Höhe 3,10 m. Der Hof wurde mit 7,20 m mal 9,00 m gemessen. Während unseres Besuchs standen dort zwei Rollcontainer. Eine Wendemöglichkeit für jeden Fahrzeugtyp wurde nicht nachgewiesen. Frau Hölscher erklärte, kleinere Lieferwagen führen rückwärts hinein und vorwärts wieder hinaus.",
        "Um 08:47 Uhr fuhr Deniz Arslan den betriebseigenen Transporter von Osten ein. Fahrzeugdaten laut seiner Zulassungsunterlage: Länge 5,99 m, Breite ohne Spiegel 2,05 m, mit Spiegeln rund 2,47 m, Höhe 2,52 m. Beim zweiten Ansetzen konnte er um die südöstliche Pflanzinsel am Eschenring fahren. Herr Brüning wies außerhalb des Fahrzeugs ein. Dafür wurde kein Gehweg überfahren. Auf dem Eschenring wartete ein Pkw etwa 35 Sekunden; ein Bus kam währenddessen nicht.",
        "Der von Getränke Röttger benannte Lkw mit 7,20 m Länge, 2,55 m Breite ohne Spiegel und 3,35 m Höhe war nicht vor Ort. Frau Hölscher erklärte, dessen Fahrer lehne den Ostweg ab. Daraus folgt aus diesem Termin weder ein erfolgreicher noch ein gescheiterter Fahrversuch dieses Lkw. Auch im Westen könnte er wegen der Torhöhe nicht in den Hof fahren. Die Skizze bildet den Ort ab, enthält aber keine berechneten Schleppkurven.",
        "Der Pfosten ließ sich nach Entlastung mit beiden Händen herausheben. Herr Arslan gab an, am 31. August habe sich der Schlüssel schwer drehen lassen. Heute gelang das Drehen ohne Werkzeug. Frau Hölscher bat darum, ihren Hinweis festzuhalten, dass eine Einweisung mit zweiter Person nicht bei jeder Lieferung möglich sei. Herr Brüning möchte die südliche Gehwegfläche frei halten; er hat einer Torverbreiterung nicht zugestimmt."
    ], "Tobias Korte\nStraßenbetrieb\nNiederschrift am 09.09.2026 um 12:10 Uhr abgeschlossen", "Ortsskizze LB-18-M1 vom 09.09.2026.")

    docx_letter("07_feuerwehr_oeffnung_2026-09-10.docx",
        ("Stadt Münster", "Feuerwehr | Vorbeugender Brandschutz | Lars Feldkamp", "lars.feldkamp@feuerwehr-muenster.example.org"),
        "Stadt Münster\nStraßenverkehr\nFrau Johanna Rensing", "10.09.2026",
        "Öffnungsprobe am Poller P1", f"FW-LB-26-73 | Bezug {AZ}", [
        "Sehr geehrte Frau Rensing,",
        "am 10. September von 07:42 bis 07:56 Uhr haben Nora Wessels und ich die Schließung des Mittelpfostens P1 geprüft. Der mitgeführte Einsatzschlüssel ließ sich verwenden. Nach dem Entriegeln wurde der Pfosten von einer Person gezogen und in die südliche Halterung gestellt. Vom Anhalten bis zur freien Durchfahrt vergingen bei dieser Probe 54 Sekunden. Der Pfosten wurde anschließend wieder eingesetzt und die Verriegelung kontrolliert.",
        "Diese Zeit ist keine Zusage für eine Einsatzfahrt. Wir waren über den Standort informiert, es herrschte Tageslicht, und unmittelbar vor dem Pfosten stand kein Fahrzeug. Bei Dunkelheit, Verschmutzung der Hülse oder einem direkt davor abgestellten Pkw kann sich der Vorgang anders darstellen. Eine vollständige Befahrung mit einem Löschfahrzeug fand nicht statt. Die gemeldete freie Breite von 3,20 m wurde von uns nicht als umfassender Nachweis für alle Einsatzfahrzeuge bewertet.",
        "Für die Objekte östlich des Pollers soll die Zufahrt bei Alarm grundsätzlich auch vom Eschenring bedacht werden. Ob dort im konkreten Moment genügend Raum vorhanden ist, hängt unter anderem von parkenden Fahrzeugen und dem Verkehr an der Einmündung ab. Die enge Kurve haben wir bei dieser Öffnungsprobe nicht mit einem Einsatzfahrzeug befahren. Eine Aussage zu notwendigen Aufstellflächen auf dem privaten Hof können wir anhand des vorliegenden Plans nicht treffen.",
        "Aus betrieblicher Sicht muss die Einsatzöffnung unabhängig von den Ladenöffnungszeiten funktionieren. Ein alleiniger Schlüssel bei Frau Hölscher wäre dafür nicht ausreichend. Bitte melden Sie uns eine Änderung der Schließanlage vor dem Austausch. Der für den Einsatz vorgesehene Zugang darf weder durch eine private Zusatzkette noch durch eine andere Sicherung ersetzt werden.",
        "Die Feuerwehr verlangt mit dieser Rückmeldung keine ganztägig offene Durchfahrt. Wir bitten um eine erneute gemeinsame Probe, falls Standort oder Pollertyp geändert werden. Zur Ausgestaltung der gewerblichen Lieferzeiten nehmen wir nicht Stellung."
    ], "Mit freundlichen Grüßen\nLars Feldkamp\nVorbeugender Brandschutz", "Keine.")

    docx_letter("08_eigentuemerin_2026-09-17.docx", OWNER,
        "Stadt Münster\nStraßenverkehr\nFrau Johanna Rensing\nVerwaltungsstelle Lindenplatz 4\n48147 Münster", "17.09.2026",
        "Zugang zum Grundstück Lindenbogen 18", f"LB18-26-09 | Ihr Vorgang {AZ}", [
        "Sehr geehrte Frau Rensing,",
        "meine Mutter Elisabeth Brüning hat mir Ihr Anhörungsschreiben an die Mieterin überlassen. Wir möchten, dass das Tor weiterhin für die Versorgung des Grundstücks erreichbar bleibt. Der Hof wird nicht erst seit dem Ladenbetrieb benutzt. Dort stehen auch die Müllbehälter für die beiden Wohnungen; gelegentlich müssen Handwerksfahrzeuge heranfahren. Wir unterstützen daher eine verlässliche Öffnungsmöglichkeit, ohne damit ein dauerndes Offenstehen des Pollers zu verlangen.",
        "Ich war beim Ortstermin am 9. September anwesend und habe Herrn Arslan eingewiesen. Sein Transporter gelangte von Osten hinein. Die Formulierung, der Ostweg sei damit allgemein geeignet, würde meine Beobachtung aber nicht vollständig wiedergeben. Es brauchte zwei Ansätze und meine Einweisung. Im Hof war an diesem Morgen außer zwei Rollcontainern nichts abgestellt. An anderen Tagen nutzen die Hausbewohner den Hof für ihre Fahrräder.",
        "Das Tor und die niedrige Durchfahrt gehören zum vorhandenen Gebäude. Wir haben Frau Hölscher keine Verbreiterung zugesagt und planen sie derzeit nicht. Auch die drei Stufen zwischen Laden und Lager werden nicht kurzfristig umgebaut. Die Hofnutzung ist mit 1.180,00 EUR Netto-Kaltmiete und 220,00 EUR Betriebskostenvorauszahlung abgegolten. Eine Mietreduzierung ist weder vereinbart noch erklärt; Frau Hölscher hat die Septemberzahlung vollständig geleistet.",
        "Vor dem Hoftor wurden in den letzten Wochen mehrfach Fahrzeuge abgestellt. Bei einem Wagen am 2. September habe ich den Fahrer aus dem Café Kranich kommen sehen. Andere Personen gingen zum Lindenhof Vorrat. Nicht jedes Auto ist einem bestimmten Betrieb zuzuordnen. Ich möchte deshalb vermeiden, dass aus den parkenden Fahrzeugen pauschal eine Kundenparkfläche des Ladens gemacht wird.",
        "Bitte beteiligen Sie uns an der Besprechung am 6. Oktober. Meine Mutter kann aus gesundheitlichen Gründen nicht selbst kommen; ich nehme für sie teil. Eine Einigung über eine neue Grundstückszufahrt oder einen Umbau des öffentlichen Gehwegs enthält dieses Schreiben nicht."
    ], "Mit freundlichen Grüßen\nNils Brüning\nfür Elisabeth Brüning", "Keine; Mietnachtrag vom 18.04.2024 liegt Frau Hölscher vor.")

    docx_letter("09_hoelscher_ergaenzung_2026-09-21.docx", SHOP,
        "Stadt Münster\nStraßenverkehr\nFrau Johanna Rensing\nVerwaltungsstelle Lindenplatz 4\n48147 Münster", "21.09.2026",
        "Ergänzung zur Anhörung über die Lieferzufahrt", f"{AZ} | Ihr Schreiben vom 14.09.2026", [
        "Sehr geehrte Frau Rensing,",
        "ich halte an meinem Antrag auf eine Öffnungsmöglichkeit von 06:00 bis 18:00 Uhr und einen zweiten Schlüssel fest. Anbei übersende ich das Fahrtenbuch für die ersten zwölf Betriebstage. Es enthält 24 Fahrten. Die darin vermerkten 192,00 EUR sind Netto-Zusatzentgelte, nicht entgangener Umsatz: 144,00 EUR von Frischefahrt West sind inzwischen berechnet, weitere 48,00 EUR von Getränke Röttger bisher nur angekündigt. Die Rechnung ist noch nicht bezahlt.",
        "Frischefahrt West kommt auf einer Tour mit mehreren Läden. Achtmal mussten wir nach 10:00 Uhr am Lindenplatz übernehmen und die Ware mit dem Handwagen etwa 65 m bis zum Hoftor bringen. Die zusätzliche Pauschale von jeweils 18,00 EUR fällt nach der Mitteilung des Unternehmens wegen der Übergabe außerhalb unserer vereinbarten Hofannahme an. Dass daraus täglich derselbe Aufwand entsteht, behaupte ich nicht. Wartezeiten in unserer Liste beruhen auf meinen oder Herrn Arslans Uhrzeiten, nicht auf einem geeichten Fahrtschreiber.",
        "Am 4. September haben wir das sofortige Wiedereinsetzen versäumt. Herr Arslan ging zu den Leerkisten, während ich allein an der Kasse stand. Unser Chat hält fest, dass der Pfosten erst um 10:19 Uhr wieder stand. Ich bestreite nicht, dass zwischenzeitlich zwei andere Pkw durchfuhren. Ich weiß aber nicht, zu wem sie gehörten. Den Schlüssel hatten wir nicht an diese Fahrer ausgegeben. Für die Zukunft kann Deniz vor 13:00 Uhr und ich anschließend selbst öffnen; während einer Bedienung bleibt trotzdem eine kurze Wartezeit möglich.",
        "Das Einstellen eines zweiten Mitarbeiters allein für den Pfosten ist für den Laden nicht tragbar. Ein zweiter Schlüssel wäre für die Überschneidung von Wareneingang und Verkauf hilfreich; er soll nicht beim Lieferdienst bleiben. Ich bin bereit, jede Öffnung mit Namen und Uhrzeit einzutragen. Ob das von Ihnen erwartete ständige Danebenstehen während der gesamten Warenannahme gemeint ist, bitte ich noch zu erläutern.",
        "Die Beschwerden von Frau Aydin nehme ich ernst. Wir haben die Kisten am 26. August nach ihrem Hinweis vom Gehweg in den Hof geräumt. Kundenfahrzeuge vor dem Tor kann ich nicht selbst entfernen. Ein generelles Versprechen, es würden keine fremden Fahrzeuge mehr halten, kann ich deshalb nicht abgeben. Am 6. Oktober kommen Herr Arslan und ich zur Besprechung."
    ], "Mit freundlichen Grüßen\nMara Hölscher\nInhaberin", "Fahrtenbuch; Schlüsseljournal; Rechnung FW-260914-118; Chat-Auszug vom 04.09.2026.")

    docx_letter("10_nachbarin_wegner_2026-09-11.docx",
        ("Ute Wegner", "Lindenbogen 16 | 48147 Münster", "ute.wegner@post.example.org"),
        "Stadt Münster\nStraßenverkehr\nFrau Johanna Rensing", "11.09.2026",
        "Beobachtungen vor den Häusern 16 und 18", f"Mein Schreiben UW-11/09 | Bezug {AZ}", [
        "Sehr geehrte Frau Rensing,",
        "ich wohne im ersten Stock des Hauses 16 und möchte meine Beobachtungen zu der Polleranlage mitteilen. Seit dem Einbau ist die Zahl der durchfahrenden Autos nach meinem Eindruck geringer. Für die zwei Wochen davor habe ich allerdings keine Zählung. Meine beigefügten Einträge betreffen einzelne Zeitfenster, in denen ich zufällig am Fenster oder vor dem Haus war; sie sind keine ganztägige Verkehrserhebung.",
        "Am Freitag, 4. September, sah ich zwischen 10:12 und 10:18 Uhr zwei Pkw von Westen kommen. Der Mittelpfosten war nicht eingesetzt. Beide fuhren am Haus vorbei nach Osten. Ich konnte nicht sehen, ob sie später am Eschenring abbogen oder an einem der Häuser anhielten. Ich habe nicht behauptet, dass ihre Fahrer Kunden von Frau Hölscher waren. Um 10:20 Uhr sah ich den Pfosten wieder stehen.",
        "Am 2. September stand von 11:03 bis mindestens 11:19 Uhr ein dunkler Kombi teilweise vor dem Hoftor. Die freie Fahrbahnbreite habe ich mit meinem Bandmaß auf etwa 2,65 m geschätzt. Ich konnte den Streifen nicht auf der ganzen Fahrzeuglänge messen. Herr Brüning sprach den Fahrer an. Ob und wann eine städtische Kontrolle kam, weiß ich nicht. Am 7. September stellte auch ein Lieferwagen den Zugang kurz zu; eine Lieferung ist für mich nicht automatisch unproblematisch.",
        "Ich habe nichts gegen den Laden oder seine Anlieferungen. Ich möchte aber nicht, dass der Pfosten jeden Vormittag herausgenommen bleibt. Meine Enkelin fährt mit dem Rad zur Schule. Wenn geparkt wird, muss sie in die Fahrbahnmitte ausweichen. Eine offene Passage hilft ihr nicht, wenn zusätzlich Autos durchfahren.",
        "Die Tabelle enthält auch Beobachtungen von Herrn Arslan und Frau Aydin. Diese Angaben habe ich nicht selbst nachgemessen; Frau Hölscher hat die Einträge für die gemeinsame Weitergabe zusammengeführt und die Namen beibehalten. Bitte behandeln Sie die unterschiedlichen Urheber entsprechend. Ich kann am 6. Oktober ab 09:30 Uhr zur Besprechung hinzukommen."
    ], "Mit freundlichen Grüßen\nUte Wegner", "Beobachtungsliste mit Einträgen bis zum 10.09.2026.")

    docx_letter("11_serviceangebot_poller_2026-09-18.docx",
        ("Klenke Metallservice", "Jasper Klenke | Werkstatt Mühlenstiege 11 | 48147 Münster", "werkstatt@klenke-metall.example.org"),
        "Stadt Münster\nStraßenbetrieb\nHerrn Tobias Korte", "18.09.2026",
        "Angebot zur Wartung des Mittelpfostens P1", f"Angebot KM-260918-09 | Ihre Anfrage vom 16.09.2026", [
        "Sehr geehrter Herr Korte,",
        "nach Besichtigung der Anlage am 17. September bieten wir die Reinigung der Bodenhülse, das Nachstellen der Verriegelung und eine gemeinsame Funktionskontrolle zum Pauschalpreis von 168,00 EUR netto an. Hinzu kommen 31,92 EUR Umsatzsteuer; der Gesamtbetrag beträgt 199,92 EUR brutto. Ersatzteile außerhalb von Kleinteilen sind darin nicht enthalten und würden vor Ausführung gesondert angeboten.",
        "Der vorhandene Mittelpfosten wiegt nach Herstelleraufkleber 8,4 kg. Wir haben ihn nicht auf einer Waage kontrolliert. Bei der Besichtigung lagen feiner Sand und zwei kleine Splittstücke in der Hülse. Nach Entlastung ließ sich der Schlüssel drehen. Ob die von Ihrem Mitarbeiter berichtete Schwergängigkeit am 31. August hierauf beruhte, können wir im Nachhinein nicht feststellen.",
        "Die seitliche Halterung war befestigt und nutzbar. Ein darin aufrecht abgestellter Pfosten ragt nicht quer über den südlichen Gehweg. Eine taktile Änderung am Gehweg oder eine Umrüstung auf einen versenkbaren Pfosten ist nicht Bestandteil dieses Angebots. Dafür wären Lage und Leitungen gesondert zu prüfen. Wir haben weder den Einsatzweg der Feuerwehr noch eine Schleppkurve für Busse oder Lieferfahrzeuge beurteilt.",
        "Bei Beauftragung bis zum 2. Oktober 2026 könnten wir die Arbeiten am 8. Oktober zwischen 07:00 und 08:00 Uhr ausführen. Der Zugang zur Hülse müsste dann frei sein. Eine Person des Straßenbetriebs soll bei der abschließenden Funktionskontrolle anwesend sein. Die Lieferung eines weiteren Schlüssels wird nicht angeboten; darüber entscheidet Ihre Schlüsselverwaltung.",
        "Dieses Angebot ist bis zum 2. Oktober 2026 gültig. Es ist noch kein Auftrag eingegangen. Die Zahlung wäre innerhalb von 14 Tagen nach Ausführung und Rechnungseingang fällig. Bitte nennen Sie bei einer Bestellung die Angebotsnummer und Ihre Rechnungsanschrift."
    ], "Mit freundlichen Grüßen\nJasper Klenke\nKlenke Metallservice", "Keine.")


def write_eml(name, sender, recipient, date, subject, body, message_id, reply_to=None, attachment=None):
    msg = EmailMessage(policy=policy.SMTPUTF8)
    msg["From"], msg["To"], msg["Date"] = contact_text(sender), contact_text(recipient), date
    msg["Subject"], msg["Message-ID"] = subject, f"<{message_id}@korrespondenz-lindenbogen-ms.de>"
    if reply_to:
        msg["In-Reply-To"] = msg["References"] = f"<{reply_to}@korrespondenz-lindenbogen-ms.de>"
    msg.set_content(contact_text(body.strip()) + "\n", charset="utf-8", cte="8bit")
    if attachment:
        path = CASE / attachment
        msg.add_attachment(path.read_bytes(), maintype="application", subtype="pdf", filename=path.name)
        msg.set_boundary("lindenhof-" + message_id)
    remember(name, subject).write_bytes(msg.as_bytes())


def build_emails():
    write_eml("12_frischefahrt_tour_2026-08-26.eml",
        "Rüdiger Maßmann <dispo@frischefahrt-west.example.org>", "Mara Hölscher <mara.hoelscher@lindenhof-vorrat.example.org>",
        "Wed, 26 Aug 2026 15:32:00 +0200", "Hofannahme Lindenbogen / Tour 4 und Zeitfenster", """
Guten Tag Frau Hölscher,

wir können Ihre Adresse nicht dauerhaft vor 10 Uhr erreichen. Tour 4 beginnt im Lager um 05:40 Uhr und bedient vor Ihnen vier Abnahmestellen mit Kühlannahme. Gestern waren wir um 10:18 Uhr bei Ihnen, heute um 10:26 Uhr. Das sind die Zeiten, die Herr Arslan am Wagen notiert hat. Unser Fahrer war heute André Lüke.

Der Wagen ist ein 3,5-t-Kühltransporter, 6,36 m lang und 2,72 m hoch. Mit Spiegeln messen wir 2,49 m Breite. Die Hofhöhe reicht dafür. Die enge Einfahrt vom Eschenring hat unser Fahrer nicht ausprobiert; er hatte niemanden zum Einweisen. Ich kann deshalb nicht bestätigen, dass dieser Weg technisch unmöglich ist.

Bis auf Weiteres übernehmen Sie verspätete Lieferungen an der Ladebucht Lindenplatz, auf Höhe Station 31 m. Bis zum Hoftor sind es nach Ihrer Messung rund 65 m mit dem Handwagen. Für diese von der vereinbarten Hofannahme abweichende Übergabe berechnen wir je tatsächlich betroffener Anlieferung 18,00 EUR netto zusätzlich. Das ist keine tägliche Pauschale. Erfolgt die Hofannahme, fällt der Zuschlag nicht an.

Bitte bestätigen Sie uns nach jeder Tour den Übergabeort. Wir fassen die Zuschläge mit den Fahrtennummern auf einer Rechnung zusammen. Einen Schlüssel Ihrer Stadt können unsere wechselnden Fahrer nicht verantwortlich verwalten.

Freundliche Grüße
Rüdiger Maßmann
Disposition Frischefahrt West
Lageradresse: Auensteg 9, 48147 Münster
""", "fw-20260826-1532")

    write_eml("13_zweitschluessel_antrag_2026-08-27.eml",
        "Mara Hölscher <mara.hoelscher@lindenhof-vorrat.example.org>", "Johanna Rensing <johanna.rensing@stadt-muenster.example.org>",
        "Thu, 27 Aug 2026 12:08:00 +0200", f"{AZ} / zweiter Schlüssel und längere Lieferzeiten", """
Sehr geehrte Frau Rensing,

ich bitte zusätzlich um einen zweiten Schlüssel. LB-17 liegt normalerweise in unserem Kassenschubfach. Wenn Deniz an der Warenannahme steht und ich einen Kunden bediene, muss er wegen des Schlüssels durch den ganzen Laden zurück. LB-18 könnte bei ihm bleiben; wir würden beide Ausgaben im Journal eintragen und keinen Schlüssel an Lieferanten geben.

Heute konnte Getränke Röttger noch um 09:25 Uhr durch P1. Frischefahrt kam dagegen erst um 11:07 Uhr. Wir haben deshalb wieder am Lindenplatz übernommen. Das Zeitfenster aus meinem Antrag vom 16. Juli, 06:00 bis 18:00 Uhr, ist weiterhin nötig. Mir geht es um einzelne Öffnungen, nicht um das ständige Entfernen des Pfostens.

Können wir den zweiten Schlüssel abholen, oder benötigen Sie weitere Angaben? Ich habe im Journal heute nur einen Antrag eingetragen. Wir haben tatsächlich weiterhin genau einen Straßenschlüssel. Die zwei Schlüssel unseres privaten Hoftors sind etwas anderes.

Mit freundlichen Grüßen
Mara Hölscher
Lindenhof Vorrat, Lindenbogen 18, Münster
""", "lh-20260827-1208")

    write_eml("14_busbetrieb_eschenring_2026-09-08.eml",
        "Anja Kemper <betrieb@quartierbus-muenster.example.org>", "Johanna Rensing <johanna.rensing@stadt-muenster.example.org>",
        "Tue, 08 Sep 2026 13:46:00 +0200", f"{AZ} / Rückmeldung Haltestelle Eschenring", """
Guten Tag Frau Rensing,

unsere Quartierbuslinie Q7 fährt auf dem Eschenring. Sie fährt weder im Regelbetrieb noch wegen des neuen Pollers durch den Lindenbogen. Bitte berichtigen Sie dies, falls Ihre Gesprächsnotiz vom 3. September einen anderen Eindruck vermittelt.

Am 2. September gegen 11:09 Uhr meldete Fahrerin Elke Struck einen wartenden Lieferwagen nahe der Einmündung. Sie konnte die Haltestelle Eschenring erst nach ungefähr zwei Minuten gerade anfahren. Wir haben keine Vermessung des Abstands zur Haltestellenkante und kein Kennzeichen notiert. Eine Zuordnung zum Lindenhof Vorrat ist damit nicht belegt.

Unser eingesetzter Solobus ist 12,00 m lang und 2,55 m breit, jeweils ohne Spiegelangabe. Für uns kommt es auf eine freie Anfahrt entlang des Eschenrings an. Eine Ersatz-Lieferstelle unmittelbar vor der Haltestelle halten wir betrieblich für ungeeignet. Ob Lieferwagen über die östliche Ecke in den Lindenbogen einbiegen können, haben wir nicht untersucht.

Wir können bei der Besprechung am 6. Oktober vertreten sein. Einen Probebus können wir an diesem Morgen jedoch nicht bereitstellen. Aus dieser Rückmeldung ergibt sich keine Zustimmung zu einer bestimmten Pollerlösung.

Freundliche Grüße
Anja Kemper
Quartierbus Münster Betriebsgesellschaft
Betriebshof Riedufer 3, 48147 Münster
""", "qb-20260908-1346")

    write_eml("15_aydin_gehweg_2026-09-10.eml",
        "Selma Aydin <selma.aydin@post.example.org>", "Tobias Korte <tobias.korte@stadt-muenster.example.org>",
        "Thu, 10 Sep 2026 18:24:00 +0200", "Lindenbogen / südlicher Gehweg und abgelegter Pfosten", """
Guten Tag Herr Korte,

ich wohne Lindenbogen 22 und fahre mit einem Aktivrollstuhl regelmäßig zum Lindenplatz. Den mittleren Poller möchte ich nicht einfach entfernt wissen, weil sonst wieder mehr Autos durchfahren. Mein Problem ist das Ablegen und die zeitweise zugestellte Gehwegfläche.

Am 26. August gegen 08:14 Uhr standen Kisten am südlichen Gehweg vor Nummer 18. Ich kam nicht gerade daran vorbei. Herr Arslan hat sie nach meiner Bitte in den Hof gestellt. Wir haben danach mit seinem Zollstock noch etwa 0,95 m freie Breite an der engsten Stelle gemessen, bevor die letzten Kisten verschwanden. Der Gehweg selbst ist deutlich breiter.

Am 4. September um 09:34 Uhr lag der herausgenommene Pfosten schräg neben der Halterung. Ich musste außen herum. Ob er dort bis 10:19 Uhr liegen blieb, weiß ich nicht; zu dieser Zeit war ich schon weg. Die beiden Ereignisse sollten nicht zu einem einzigen Foto oder einer dauernden Blockade zusammengezogen werden.

Ich brauche eine Lösung, bei der der Gehweg tatsächlich frei bleibt und die Pfosten im Dunkeln gut erkennbar sind. Aus dem Rollstuhl kann ich keinen schweren Mittelpfosten heben. Eine allgemeine Beurteilung der Anlage nach technischen Normen kann ich nicht abgeben. Bitte nehmen Sie mich zur Ortsbesichtigung mit, falls die Halterung versetzt werden soll.

Mit freundlichen Grüßen
Selma Aydin
Lindenbogen 22, 48147 Münster
""", "sa-20260910-1824")

    write_eml("16_rensing_portalnachricht_2026-09-15.eml",
        "Johanna Rensing <johanna.rensing@stadt-muenster.example.org>", "Mara Hölscher <mara.hoelscher@lindenhof-vorrat.example.org>",
        "Tue, 15 Sep 2026 11:06:00 +0200", f"{AZ} / Nachricht zur Schlüsselverwaltung", """
Sehr geehrte Frau Hölscher,

Ihre Nachricht ist eingegangen. Eine Entscheidung über den zweiten Schlüssel liegt noch nicht vor. Die bisherige Regelung bleibt bis 15.10.2026 unverändert. Bitte reichen Sie Ihre Ergänzung zur Anhörung bis 30.09.2026 ein.

Der Eintrag LB-18 in der Schlüsselverwaltung bezeichnet den noch offenen Zusatzantrag vom 27. August. Er ist keine Ausgabebestätigung. Bitte holen Sie ohne weitere Nachricht keinen Schlüssel im Werkhof ab. Die ausgegebene Kennung lautet weiterhin LB-17; Herr Arslan ist als Ihre betriebliche Vertretung benannt.

Der Termin am 6. Oktober um 09:00 Uhr ist für Sie und Herrn Arslan vorgemerkt. Ich habe Herrn Brüning um eine gesonderte Rückmeldung gebeten. Wenn Ihre Fahrtenliste bislang nur bis zum 8. September reicht, können Sie diese schon jetzt einreichen und spätere Tage nachreichen. Eine Verlängerung der schriftlichen Frist ist damit nicht verbunden.

Mit freundlichen Grüßen
Johanna Rensing
Stadt Münster, Straßenverkehr
Verwaltungsstelle Lindenplatz 4, 48147 Münster
""", "ms-20260915-1106", reply_to="lh-20260827-1208")

    write_eml("17_frischefahrt_rechnung_2026-09-14.eml",
        "Rüdiger Maßmann <dispo@frischefahrt-west.example.org>", "Mara Hölscher <mara.hoelscher@lindenhof-vorrat.example.org>",
        "Mon, 14 Sep 2026 09:18:00 +0200", "FW-260914-118 / Zusatzübergaben am Lindenplatz", """
Guten Tag Frau Hölscher,

anbei die gesonderte Rechnung FW-260914-118 über acht Übergaben am Lindenplatz. Betroffen sind der 25., 26., 27., 28. und 31. August sowie der 1., 2. und 7. September. Unsere Fahrtenzuordnung entspricht Ihren Einträgen F03, F06, F08, F09, F11, F14, F15 und F21. Die Ware selbst ist nicht auf dieser Rechnung abgerechnet.

Der Zuschlag beträgt jeweils 18,00 EUR netto, zusammen 144,00 EUR zuzüglich 27,36 EUR Umsatzsteuer. Bitte zahlen Sie die 171,36 EUR bis 28. September auf das Ihnen aus den Warenrechnungen bekannte Geschäftskonto. Es handelt sich nicht um bereits eingezogene Beträge. Eine Gutschrift haben wir bisher nicht erstellt.

Am 4. September war die Hofübergabe möglich, deshalb erscheint diese Fahrt nicht auf der Rechnung. Dass Herr Lüke vor der Einfahrt warten musste, löst nach unserer Absprache allein keinen Zuschlag aus. Die Zuschläge von Getränke Röttger sind nicht unsere Forderungen und dürfen nicht nochmals zu unserem Rechnungsbetrag addiert werden.

Freundliche Grüße
Rüdiger Maßmann
Frischefahrt West, Auensteg 9, 48147 Münster
""", "fw-20260914-0918", reply_to="fw-20260826-1532", attachment="26_rechnung_frischefahrt_2026-09-14.pdf")

    write_eml("18_cafe_kundenparken_2026-09-02.eml",
        "Bettina Seidel <kontakt@cafe-kranich.example.org>", "Mara Hölscher <mara.hoelscher@lindenhof-vorrat.example.org>",
        "Wed, 02 Sep 2026 16:11:00 +0200", "Wagen vor eurem Tor heute Vormittag", """
Hallo Mara,

der Fahrer des dunklen Kombis war heute bei uns im Café Kranich, Lindenbogen 14. Herr Brüning hat mich gegen 11:15 Uhr angesprochen, und ich habe im Gastraum nachgefragt. Der Mann ist dann rausgegangen. Ich habe nicht kontrolliert, ob er sofort weggefahren ist. Einen Namen oder das Kennzeichen habe ich nicht.

Ich habe ihm gesagt, dass er nicht vor eurem Tor stehen darf. Er meinte, er sei nur zum Abholen da und habe niemanden behindert. Das war nicht unsere Erlaubnis, und wir haben keine Kundenstellplätze auf der Straße. Wir können unsere Gäste darauf ansprechen, aber ich kann nicht jeden Ankommenden draußen überwachen.

Bitte ordnet uns nicht sämtliche Autos zu, die dort stehen. Gestern hat jemand aus eurem Laden Tüten in einen anderen Wagen getragen. Umgekehrt stimmt also auch nicht, dass alle Wagen von uns kommen. Wir sollten beide dafür sorgen, dass Abholer die Zufahrt nicht als Kurzzeitparkplatz ansehen.

Viele Grüße
Bettina
Bettina Seidel, Café Kranich
Lindenbogen 14, 48147 Münster
""", "ck-20260902-1611")

    write_eml("19_eingang_aktenauskunft_2026-09-22.eml",
        "Johanna Rensing <johanna.rensing@stadt-muenster.example.org>", "Mara Hölscher <mara.hoelscher@lindenhof-vorrat.example.org>",
        "Tue, 22 Sep 2026 14:37:00 +0200", f"{AZ} / Eingang Ihrer Ergänzung und Unterlagenstand", """
Sehr geehrte Frau Hölscher,

Ihre Ergänzung vom 21. September ist mit vier Anlagen eingegangen. Die Frist bis zum 30. September bleibt für weitere Unterlagen offen. Eine inhaltliche Entscheidung ist heute nicht getroffen worden.

Herr Brüning hat zusätzlich nach der früheren straßenrechtlichen Widmung und den Unterlagen zur vorhandenen Grundstückszufahrt gefragt. Aus der laufenden Verkehrsakte kann ich dazu bislang nur mitteilen, dass der Lindenbogen als kommunale öffentliche Straße geführt wird. Die ursprüngliche Widmungsverfügung und die ältere Bauakte des Tors sind in der Ihnen überlassenen Teilakte nicht enthalten. Unsere Registratur sucht diese Unterlagen; ich sage Ihnen noch keinen Fund oder bestimmten Inhalt zu.

Die verkehrsrechtliche Anordnung SV-LB-26-118 ist in unserer Mitteilung vom 24. Juli bezeichnet. Die Mitteilung ersetzt nicht die vollständige Anordnung mit Anlagen. Auch deren vollständige Kopie ist bei der Registratur angefordert. Wir werden fehlende Unterlagen nachreichen, sobald sie vorliegen. Der Ortsterminplan vom 9. September ist eine neue Messskizze und kein Katasterauszug.

Mit freundlichen Grüßen
Johanna Rensing
Stadt Münster, Straßenverkehr
""", "ms-20260922-1437")


TRIPS = [
    ("F01", "2026-08-24", "07:12", "Frischefahrt West", "André Lüke", "Kühltransporter 3,5 t", "West / Hof", 8, "0,00", "Lieferschein FW-8241", "Öffnung durch Arslan; Hofannahme"),
    ("F02", "2026-08-24", "10:42", "Getränke Röttger", "Jens Rößler", "Lkw 7,5 t", "Lindenplatz", 18, "12,00", "Ankündigung GR-24", "Handwagen ab Ladebucht; Zuschlag angekündigt"),
    ("F03", "2026-08-25", "10:18", "Frischefahrt West", "André Lüke", "Kühltransporter 3,5 t", "Lindenplatz", 17, "18,00", "FW-260914-118 / 1", "Übergabe außerhalb Hofannahme"),
    ("F04", "2026-08-25", "13:06", "Paketdienst Auenkurier", "Nina Voß", "Kastenwagen 2,8 t", "Ost / Hof", 0, "0,00", "Empfang AK-825", "Zwei Kartons; kurze Einfahrt ohne Einweiser"),
    ("F05", "2026-08-26", "08:06", "Gemüsehof Berken", "Paul Öster", "Transporter 3,5 t", "West / Hof", 5, "0,00", "Lieferschein GB-826", "Kisten kurz auf Gehweg; nach Hinweis weggeräumt"),
    ("F06", "2026-08-26", "10:26", "Frischefahrt West", "André Lüke", "Kühltransporter 3,5 t", "Lindenplatz", 19, "18,00", "FW-260914-118 / 2", "Tourzeit von Arslan notiert"),
    ("F07", "2026-08-27", "09:25", "Getränke Röttger", "Jens Rößler", "Lkw 7,5 t", "West / vor Tor", 8, "0,00", "Lieferschein GR-827", "Zu hoch für Hof; Straßenentladung elf Minuten"),
    ("F08", "2026-08-27", "11:07", "Frischefahrt West", "Heike Möller", "Kühltransporter 3,5 t", "Lindenplatz", 22, "18,00", "FW-260914-118 / 3", "Hölscher zunächst allein im Verkauf"),
    ("F09", "2026-08-28", "10:14", "Frischefahrt West", "André Lüke", "Kühltransporter 3,5 t", "Lindenplatz", 16, "18,00", "FW-260914-118 / 4", "Vier Rollkisten mit eigenem Handwagen übernommen"),
    ("F10", "2026-08-28", "14:22", "Paketdienst Auenkurier", "Nina Voß", "Kastenwagen 2,8 t", "Ost / Hof", 0, "0,00", "Empfang AK-828", "Drei Kartons; keine Schlüsselöffnung"),
    ("F11", "2026-08-31", "10:34", "Frischefahrt West", "Heike Möller", "Kühltransporter 3,5 t", "Lindenplatz", 18, "18,00", "FW-260914-118 / 5", "Keine Öffnung nach 10 Uhr"),
    ("F12", "2026-08-31", "12:16", "Getränke Röttger", "Jens Rößler", "Lkw 7,5 t", "Lindenplatz", 15, "12,00", "Ankündigung GR-31", "Zuschlag nur angekündigt; Leergebinde mitgenommen"),
    ("F13", "2026-09-01", "08:15", "Gemüsehof Berken", "Paul Öster", "Transporter 3,5 t", "West / Hof", 3, "0,00", "Lieferschein GB-901", "Arslan wartet an P1"),
    ("F14", "2026-09-01", "10:48", "Frischefahrt West", "André Lüke", "Kühltransporter 3,5 t", "Lindenplatz", 21, "18,00", "FW-260914-118 / 6", "Wartezeit bis Beginn der Übergabe"),
    ("F15", "2026-09-02", "11:12", "Frischefahrt West", "Heike Möller", "Kühltransporter 3,5 t", "Lindenplatz", 24, "18,00", "FW-260914-118 / 7", "Kombi vor Hoftor; Handwagen musste kurz warten"),
    ("F16", "2026-09-02", "13:38", "Paketdienst Auenkurier", "Nina Voß", "Kastenwagen 2,8 t", "Ost / Hof", 0, "0,00", "Empfang AK-902", "Hoftor wieder frei"),
    ("F17", "2026-09-03", "07:58", "Gemüsehof Berken", "Paul Öster", "Transporter 3,5 t", "West / Hof", 4, "0,00", "Lieferschein GB-903", "Öffnung und Schließung durch Arslan"),
    ("F18", "2026-09-03", "11:35", "Getränke Röttger", "Jens Rößler", "Lkw 7,5 t", "Lindenplatz", 20, "12,00", "Ankündigung GR-03", "Zwei Handwagenwege; Zuschlag angekündigt"),
    ("F19", "2026-09-04", "09:26", "Frischefahrt West", "André Lüke", "Kühltransporter 3,5 t", "West / Hof", 12, "0,00", "Lieferschein FW-904", "Durchfahrt um 09:38; späteres Schließen laut Chat 10:19"),
    ("F20", "2026-09-04", "10:58", "Getränke Röttger", "Jens Rößler", "Lkw 7,5 t", "Lindenplatz", 17, "12,00", "Ankündigung GR-04", "P1 bereits wieder eingesetzt; Zuschlag angekündigt"),
    ("F21", "2026-09-07", "10:21", "Frischefahrt West", "Heike Möller", "Kühltransporter 3,5 t", "Lindenplatz", 23, "18,00", "FW-260914-118 / 8", "Übergabe am Lindenplatz; Zuordnung fremder Lieferwagen offen"),
    ("F22", "2026-09-07", "14:04", "Paketdienst Auenkurier", "Nina Voß", "Kastenwagen 2,8 t", "Ost / Hof", 0, "0,00", "Empfang AK-907", "Keine Polleröffnung"),
    ("F23", "2026-09-08", "08:12", "Gemüsehof Berken", "Paul Öster", "Transporter 3,5 t", "West / Hof", 4, "0,00", "Lieferschein GB-908", "Tor durch Arslan geöffnet"),
    ("F24", "2026-09-08", "09:16", "Getränke Röttger", "Jens Rößler", "Lkw 7,5 t", "West / vor Tor", 7, "0,00", "Lieferschein GR-908", "Straßenentladung neun Minuten; kein Hofeintritt"),
]

KEYS = [
    ("S01", "2026-08-21", "08:20", "LB-17", "Tobias Korte", "Mara Hölscher", "Ausgabe Werkhof", "2026-10-16 12:00", "Einweisung; nur ein Straßenschlüssel ausgegeben"),
    ("S02", "2026-08-24", "06:25", "LB-17", "Mara Hölscher", "Deniz Arslan", "Interne Übergabe", "2026-08-24 10:00", "Kassenschubfach; Frühdienst"),
    ("S03", "2026-08-24", "09:48", "LB-17", "Deniz Arslan", "Mara Hölscher", "Interne Rückgabe", "", "Pfosten verriegelt"),
    ("S04", "2026-08-25", "09:55", "LB-17", "Mara Hölscher", "Kassenschubfach", "Verwahrung", "", "Keine Öffnung für spätere Frischefahrt"),
    ("S05", "2026-08-26", "07:50", "LB-17", "Mara Hölscher", "Deniz Arslan", "Interne Übergabe", "2026-08-26 10:00", "Gemüseanlieferung"),
    ("S06", "2026-08-26", "09:40", "LB-17", "Deniz Arslan", "Kassenschubfach", "Verwahrung", "", "Rückgabe von Hölscher gesehen"),
    ("S07", "2026-08-27", "12:08", "LB-18", "Mara Hölscher", "Johanna Rensing", "Zusatzantrag", "", "Nicht ausgegeben; E-Mail vom selben Tag"),
    ("S08", "2026-08-31", "08:03", "LB-17", "Deniz Arslan", "Tobias Korte", "Störungsmeldung", "", "Telefonisch schwergängiges Drehen gemeldet; kein Schlüsselwechsel"),
    ("S09", "2026-09-01", "07:55", "LB-17", "Mara Hölscher", "Deniz Arslan", "Interne Übergabe", "2026-09-01 10:00", "Gemüseanlieferung"),
    ("S10", "2026-09-01", "09:45", "LB-17", "Deniz Arslan", "Kassenschubfach", "Verwahrung", "", "Pfosten verriegelt"),
    ("S11", "2026-09-04", "09:20", "LB-17", "Mara Hölscher", "Deniz Arslan", "Interne Übergabe", "2026-09-04 10:00", "Frischefahrt angekündigt"),
    ("S12", "2026-09-04", "10:19", "LB-17", "Mara Hölscher", "Kassenschubfach", "Verwahrung", "", "Schließen laut Chat; Zeitpunkt Rückgabe Arslan an Hölscher nicht notiert"),
    ("S13", "2026-09-09", "08:30", "LB-17", "Mara Hölscher", "Tobias Korte", "Funktionsprobe vor Ort", "2026-09-09 09:25", "Keine dauerhafte Rückgabe an Stadt"),
    ("S14", "2026-09-09", "09:25", "LB-17", "Tobias Korte", "Mara Hölscher", "Rückgabe nach Probe", "2026-10-16 12:00", "Derselbe Schlüssel; Versuch unverändert"),
    ("S15", "2026-09-15", "11:06", "LB-18", "Johanna Rensing", "Mara Hölscher", "Statusmitteilung", "", "In Prüfung; nicht ausgegeben"),
    ("S16", "2026-09-22", "15:42", "LB-17", "Betriebsablage", "Mara Hölscher", "Bestandsabgleich", "2026-10-16 12:00", "Ausgegeben; Versuch bis 15.10.2026"),
    ("S17", "2026-09-22", "15:42", "LB-18", "Betriebsablage", "Mara Hölscher", "Bestandsabgleich", "", "Nicht ausgegeben; Zusatzantrag in Prüfung"),
]

OBSERVATIONS = [
    ("B01", "2026-08-24", "07:12", "07:24", "P1", "Deniz Arslan", 1, "3,20", "Frischefahrt passiert; Breite bei geöffnetem Pfosten"),
    ("B02", "2026-08-24", "15:20", "15:35", "Vor Hoftor 18", "Ute Wegner", 1, "2,70", "Pkw halb vor Tor; Ziel unbekannt; Restfahrbahn geschätzt"),
    ("B03", "2026-08-26", "08:14", "08:18", "Gehweg Süd vor 18", "Selma Aydin", 0, "0,95", "Restgehweg neben Kisten; danach freigeräumt"),
    ("B04", "2026-08-27", "09:33", "09:44", "Fahrbahn vor Hoftor", "Deniz Arslan", 1, "2,25", "Getränke-Lkw entlädt; rechnerisch 4,80 minus 2,55 ohne Spiegel"),
    ("B05", "2026-08-28", "16:10", "16:22", "Vor Hoftor 18", "Ute Wegner", 1, "", "Fahrer trägt Tüte aus Laden; keine Breitenmessung"),
    ("B06", "2026-08-31", "08:03", "08:07", "P1", "Deniz Arslan", 0, "", "Schlüssel ließ sich zunächst schwer drehen; keine Fahrzeugzählung"),
    ("B07", "2026-09-01", "11:25", "11:40", "Vor Haus 16", "Ute Wegner", 0, "4,80", "Kein abgestelltes Fahrzeug in diesem Abschnitt"),
    ("B08", "2026-09-02", "11:03", "11:19", "Vor Hoftor 18", "Ute Wegner", 1, "2,65", "Dunkler Kombi; Restfahrbahn nur punktuell gemessen"),
    ("B09", "2026-09-02", "11:09", "11:11", "Eschenring Haltestelle", "Elke Struck", 1, "", "Lieferwagen nahe Einmündung; betriebliche Meldung über Kemper"),
    ("B10", "2026-09-03", "17:15", "17:28", "Vor Haus 14", "Ute Wegner", 2, "", "Zwei Abholer nacheinander; Kundenzuordnung unklar"),
    ("B11", "2026-09-04", "09:34", "09:35", "P1 Gehweg Süd", "Selma Aydin", 0, "", "Pfosten lag schräg neben Halterung; keine Breitenmessung"),
    ("B12", "2026-09-04", "10:12", "10:18", "P1 Richtung Osten", "Ute Wegner", 2, "3,20", "Zwei Pkw durch offene Passage; Ziele nicht beobachtet"),
    ("B13", "2026-09-04", "10:20", "10:21", "P1", "Ute Wegner", 0, "1,50", "Mittelpfosten eingesetzt; Wert je seitlicher Öffnung"),
    ("B14", "2026-09-07", "10:18", "10:27", "Vor Hoftor 18", "Ute Wegner", 1, "2,55", "Lieferwagen; Firmenname nicht erkannt; Restfahrbahn geschätzt"),
    ("B15", "2026-09-08", "09:23", "09:32", "Fahrbahn vor Hoftor", "Deniz Arslan", 1, "2,25", "Getränkeentladung; rechnerischer Wert ohne Spiegel"),
    ("B16", "2026-09-09", "08:47", "08:49", "Einmündung Eschenring", "Tobias Korte", 1, "", "Betriebstransporter zwei Ansätze mit Einweiser; Pkw wartete"),
    ("B17", "2026-09-10", "07:42", "07:56", "P1", "Lars Feldkamp", 0, "", "Öffnungsprobe; keine Einsatzfahrzeugbefahrung"),
    ("B18", "2026-09-18", "15:05", "15:20", "Gehweg Süd vor 18", "Selma Aydin", 0, "2,00", "Keine Kisten; Gehweg durchgehend frei"),
]


def write_csv(name, description, header, rows):
    with remember(name, description).open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file, delimiter=";", lineterminator="\n")
        writer.writerow(header.split(";"))
        writer.writerows(rows)


def build_csvs():
    write_csv("23_fahrtenbuch_2026-08-24_bis_09-08.csv", "24 Fahrten mit Übergabeorten, Wartezeiten und Zusatzentgelten",
        "Fahrt_ID;Datum;Ankunft;Betrieb;Fahrer;Fahrzeug;Zugang;Wartezeit_min;Zusatzentgelt_netto_EUR;Beleg;Notiz", TRIPS)
    write_csv("24_schluesseljournal_2026-08-21_bis_09-22.csv", "17 Schlüsselvorgänge und Statusmeldungen",
        "Vorgang_ID;Datum;Uhrzeit;Schlüssel;Von;An;Vorgang;Rückgabe_bis;Notiz", KEYS)
    write_csv("25_beobachtungen_lindenbogen_2026-08-24_bis_09-18.csv", "18 örtliche Beobachtungen mit unterschiedlichen Messbezügen",
        "Beobachtung_ID;Datum;Von;Bis;Ort;Erfasser;Fahrzeuge_Anzahl;Restbreite_m;Beobachtung", OBSERVATIONS)


def build_texts():
    remember("20_chat_fruehdienst_2026-09-04.txt", "Chat-Auszug zum verspäteten Wiedereinsetzen des Pollers").write_text("""Lindenhof Frühdienst
Export durch Mara Hölscher am 21.09.2026, 18:06 Uhr
Teilnehmende: Mara Hölscher, Deniz Arslan
Zeitangaben: Ortszeit Münster; Auszug vom 04.09.2026

04.09.2026 09:38 | Deniz Arslan
Frischefahrt ist durch. Ich räume die leeren Kisten hinten weg. Kannst du den Poller wieder einsetzen?

04.09.2026 09:41 | Mara Hölscher
Bin an der Kasse, fünf Minuten. Bitte den Schlüssel nicht im Schloss lassen.

04.09.2026 10:12 | Deniz Arslan
Steht noch offen. Zwei Autos sind eben hintereinander rein, keines gehört zu unserer Lieferung.

04.09.2026 10:19 | Mara Hölscher
Jetzt geschlossen, LB-17 liegt im Schubfach. Ich habe nicht gesehen, wohin die beiden gefahren sind.

In diesem Auszug wurden keine Bilder oder Sprachnachrichten mit exportiert.
Der Ausdruck enthält nur den ausgewählten Zeitraum, nicht den gesamten Gruppenverlauf.
""", encoding="utf-8")
    remember("21_schluesselablage_2026-09-22.txt", "Interne Betriebsablage zu LB-17 und offenem Zusatzantrag LB-18").write_text("""Lindenhof Vorrat | Betriebliche Schlüsselablage
Interner Stand von Mara Hölscher am 22.09.2026 um 15:42 Uhr
Vorgang: MS-LB-2026-0417
Standort: Lindenbogen P1

Schlüssel LB-17
Empfängerin: Mara Hölscher
Übernommen: 21.08.2026 08:20 Uhr
Status: ausgegeben
Versuch: 24.08.2026 bis 15.10.2026, Montag bis Freitag 06:30 bis 10:00 Uhr
Rückgabe bis: 16.10.2026 12:00 Uhr

Schlüssel LB-18
Antragsteller: Lindenhof Vorrat
Zusatzantrag: 27.08.2026
Ausgabe: nicht ausgegeben
Status: in Prüfung

Aus E-Mail Johanna Rensing, 15.09.2026 11:06 Uhr übernommen
Ihre Nachricht ist eingegangen. Eine Entscheidung über den zweiten Schlüssel liegt noch nicht vor. Die bisherige Regelung bleibt bis 15.10.2026 unverändert. Bitte reichen Sie Ihre Ergänzung zur Anhörung bis 30.09.2026 ein.

Interne Betriebsnotiz nach Ausgabebeleg und E-Mail, keine amtliche Portalbestätigung. Private Hoftorschlüssel werden in dieser Ablage nicht geführt.
""", encoding="utf-8")
    remember("22_telefonnotiz_roettger_2026-09-18.txt", "Telefonnotiz zu Getränkelieferung, angekündigten Entgelten und Fahrzeugmaßen").write_text("""Lindenhof Vorrat | Telefonbuch Wareneingang
18.09.2026, 12:22 bis 12:34 Uhr
Gespräch: Mara Hölscher mit Jens Rößler, Getränke Röttger
Von Mara Hölscher um 12:45 Uhr niedergeschrieben, Herrn Rößler nicht zur Bestätigung vorgelegt.

Rößler sagt, die vier Zusatzentgelte für die Übergabe am Lindenplatz vom 24. und 31. August sowie 3. und 4. September betrügen jeweils 12,00 EUR netto. Es liege noch keine Rechnung vor. Die nächste Sammelrechnung solle Ende September kommen. Die beiden Lieferungen vor dem Tor am 27. August und 8. September bekämen keinen solchen Zuschlag. Die 48,00 EUR seien daher nur vorgemerkt, nicht bezahlt.

Sein regulärer Lkw habe laut Fahrzeugblatt 7,20 m Länge, 2,55 m Breite ohne Spiegel und 3,35 m Höhe. In das 3,10 m hohe Hoftor könne er nicht hineinfahren. Das sei auch vor dem Poller so gewesen. Bisher habe er von Westen bis vor das Tor fahren und dort abladen können. Den Ostweg habe er mit dem beladenen Lkw nicht ausprobiert. Er halte die Ecke an der Pflanzinsel für zu eng und wolle dort ohne gesicherte Einweisung nicht rangieren.

Ein kleinerer Wagen sei im Betrieb vorhanden, aber nicht täglich frei. Er wolle fragen, ob für Oktober eine feste Vormittagstour angeboten werden könne. Preis und Tage seien noch offen. Eine Zusage, alle Lieferungen künftig vor 10 Uhr zu bringen, habe er nicht gegeben.

Ich habe ihm erklärt, dass der zweite Straßenschlüssel bislang nicht ausgegeben ist. Er soll keinen eigenen Nachschlüssel anfertigen lassen. Für den 6. Oktober ist er auf Tour und kann an dem Gespräch bei der Stadt nicht teilnehmen. Er ist mit der Weitergabe seiner Fahrzeugangaben einverstanden; eine schriftliche Fahrbarkeitsprüfung ersetzt das Gespräch nicht.

Mara Hölscher
""", encoding="utf-8")


def pdf_paragraph(c, text, x, y, width, size=11, leading=14):
    style = ParagraphStyle("p", fontName="Times-Roman", fontSize=size, leading=leading)
    para = Paragraph(escape(contact_text(text)).replace("\n", "<br/>"), style)
    _, height = para.wrap(width, 1000)
    para.drawOn(c, x, y - height)
    return y - height - 9


def build_invoice():
    path = remember("26_rechnung_frischefahrt_2026-09-14.pdf", "Rechnung über acht Zusatzübergaben mit 171,36 EUR Bruttobetrag")
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    c.setTitle("Rechnung FW-260914-118")
    c.setAuthor("Frischefahrt West")
    x, y, width = 24*mm, 274*mm, 162*mm
    c.setFont("Times-Bold", 17)
    c.drawString(x, y, "Frischefahrt West")
    y = pdf_paragraph(c, "Rüdiger Maßmann | Auensteg 9 | 48147 Münster\ndispo@frischefahrt-west.example.org", x, y-15, width, 10, 13)
    y = pdf_paragraph(c, "Lindenhof Vorrat\nFrau Mara Hölscher\nLindenbogen 18\n48147 Münster", x, y-15, width)
    c.setFont("Times-Bold", 14)
    c.drawString(x, y-16, "Rechnung FW-260914-118")
    y = pdf_paragraph(c, "Rechnungsdatum 14.09.2026 | Kundennummer LHV-018\nLeistungszeitraum 25.08. bis 07.09.2026", x, y-28, width, 10, 13)
    y = pdf_paragraph(c, "Guten Tag Frau Hölscher,\nfür die gesonderte Übergabe am Lindenplatz berechnen wir die folgenden, mit Ihrem Fahrtenbuch abgestimmten Zuschläge. Warenwerte sind nicht enthalten.", x, y-4, width)
    columns = [x, x+16*mm, x+46*mm, x+121*mm, x+161*mm]
    c.setFillColorRGB(.93, .93, .93)
    c.rect(x, y-19, width, 20, fill=1, stroke=0)
    c.setFillColorRGB(0,0,0)
    c.setFont("Times-Bold", 10)
    for pos, label in zip(columns, ["Pos.", "Datum", "Fahrt und Leistung", "Anzahl", "Netto EUR"]):
        (c.drawRightString if label == "Netto EUR" else c.drawString)(pos, y-13, label)
    y -= 21
    charged = [row for row in TRIPS if row[8] == "18,00"]
    c.setFont("Times-Roman", 10)
    for idx, row in enumerate(charged, 1):
        y -= 23
        c.drawString(columns[0], y, str(idx))
        c.drawString(columns[1], y, datetime.fromisoformat(row[1]).strftime("%d.%m.%Y"))
        c.drawString(columns[2], y, row[0] + " Übergabe Lindenplatz")
        c.drawString(columns[3], y, "1")
        c.drawRightString(columns[4], y, "18,00")
        c.setStrokeColorRGB(.85,.85,.85)
        c.line(x, y-7, x+width, y-7)
    y -= 32
    for label, amount in [("Summe netto", "144,00"), ("Umsatzsteuer 19 %", "27,36"), ("Rechnungsbetrag EUR", "171,36")]:
        c.setFont("Times-Bold" if label.startswith("Rechnungs") else "Times-Roman", 11)
        c.drawString(x+78*mm, y, label)
        c.drawRightString(x+width, y, amount)
        y -= 19
    y = pdf_paragraph(c, "Zahlbar ohne Abzug bis 28.09.2026 auf das aus den Warenrechnungen bekannte Geschäftskonto. Bitte die Rechnungsnummer angeben. Ein Zahlungseinzug ist nicht erfolgt.", x, y-5, width)
    pdf_paragraph(c, "Vielen Dank.\nRüdiger Maßmann", x, y, width)
    c.setFont("Times-Roman", 8)
    c.drawString(x, 15*mm, "FW-260914-118 | gesonderte Transportleistung | Seite 1")
    c.save()


def build_plan():
    path = remember("27_ortsskizze_lindenbogen_2026-09-09.pdf", "Messskizze mit Pollerquerschnitt, Hof und östlichem Anschluss")
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    c.setTitle("Ortsskizze LB-18-M1")
    c.setAuthor("Tobias Korte")
    x, width = 22*mm, 166*mm
    c.setFont("Times-Bold", 15)
    c.drawString(x, 277*mm, "Lindenbogen und Grundstückszufahrt")
    pdf_paragraph(c, "Straßenbetrieb | Tobias Korte | 09.09.2026\nMS-LB-2026-0417 | Skizze LB-18-M1 | Messung mit Bandmaß", x, 270*mm, width, 10, 13)
    c.setFont("Times-Roman", 10)
    c.drawString(x, 246*mm, "Lageskizze, nicht maßstäblich. Norden oben.")
    # Längsabstände und Querschnitt haben bewusst getrennte Darstellungen.
    left, right, cy = 36*mm, 171*mm, 204*mm
    c.setFillColorRGB(.96,.96,.96)
    c.rect(left, cy-15*mm, right-left, 30*mm, fill=1, stroke=1)
    c.setFillColorRGB(0,0,0)
    c.line(left,cy-10*mm,right,cy-10*mm)
    c.line(left,cy+10*mm,right,cy+10*mm)
    c.setFont("Times-Roman", 10)
    c.drawString(left+2*mm,cy+11.5*mm,"Gehweg Nord 1,80 m")
    c.drawString(left+2*mm,cy-13.5*mm,"Gehweg Süd 2,00 m")
    c.drawString(left+55*mm,cy+2*mm,"Lindenbogen | Fahrbahn 4,80 m")
    c.drawString(left+55*mm,cy-4*mm,"Fahrradstraße / Anlieger frei")
    for pos, label in [(left,"0 m"),(left+48*mm,"64 m"),(left+72*mm,"96 m"),(right,"180 m")]:
        c.line(pos,cy+16*mm,pos,cy+19*mm)
        c.drawCentredString(pos,cy+21*mm,label)
    px, gx = left+48*mm, left+72*mm
    for py in [cy-9*mm,cy,cy+9*mm]:
        c.circle(px,py,1.3*mm,fill=1)
    c.drawString(px-5*mm,cy-23*mm,"P1")
    c.line(gx,cy-15*mm,gx,cy-22*mm)
    c.rect(gx-11*mm,cy-49*mm,38*mm,27*mm)
    c.drawCentredString(gx+8*mm,cy-29*mm,"Hof 7,20 × 9,00 m")
    c.drawCentredString(gx+8*mm,cy-35*mm,"Lindenhof Vorrat")
    c.drawCentredString(gx+8*mm,cy-41*mm,"Lindenbogen 18")
    c.drawString(left,cy-34*mm,"Tor 2,85 m breit")
    c.drawString(left,cy-40*mm,"Tor 3,10 m hoch")
    c.saveState()
    c.translate(left-8*mm, cy-10*mm)
    c.rotate(90)
    c.drawString(0,0,"Lindenplatz (West)")
    c.restoreState()
    c.line(right+2*mm,cy-27*mm,right+2*mm,cy+31*mm)
    c.drawString(right+4*mm,cy+23*mm,"Bus")
    c.drawString(right+4*mm,cy+18*mm,"Q7")
    c.saveState()
    c.translate(right+13*mm,cy-13*mm)
    c.rotate(90)
    c.drawString(0,0,"Eschenring (Ost)")
    c.restoreState()
    c.setFillColorRGB(.8,.85,.8)
    c.rect(right-8*mm,cy-15*mm,8*mm,5*mm,fill=1,stroke=1)
    c.setFillColorRGB(0,0,0)
    c.drawString(right-24*mm,cy-30*mm,"Pflanzinsel")
    c.line(right-12*mm,cy-27*mm,right-4*mm,cy-14*mm)
    pdf_paragraph(c, "Ladebucht Lindenplatz bei Station 31 m; Transportweg bis zum Hoftor rund 65 m. Abstand P1 bis Tor 32 m. Eschenring: Haltestellenanfahrt östlich des Anschlusses; keine Busroute im Lindenbogen.", x, 146*mm, width, 10, 13)
    c.setFont("Times-Bold", 12)
    c.drawString(x,120*mm,"Querschnitt P1")
    start, end, base = 49*mm, 161*mm, 105*mm
    c.line(start,base,end,base)
    for pos, size in [(start,1.4*mm),((start+end)/2,2.8*mm),(end,1.4*mm)]:
        c.setFillColorRGB(.35,.35,.35)
        c.rect(pos-size/2,base-5*mm,size,10*mm,fill=1,stroke=1)
    c.setFillColorRGB(0,0,0)
    c.setFont("Times-Roman", 10)
    c.drawCentredString((start*3+end)/4,base+8*mm,"1,50 m frei")
    c.drawCentredString((start+end*3)/4,base+8*mm,"1,50 m frei")
    c.drawCentredString((start+end)/2,base-11*mm,"Mittelpfosten 0,20 m; herausgenommen: 3,20 m frei")
    y = pdf_paragraph(c, "Die Maße beziehen sich auf sichtbare Kanten. Die Durchfahrtbreite ist kein Nachweis einer sicheren Befahrung mit jedem Fahrzeug. Spiegelbreiten, Rangieren und überhängende Fahrzeugteile sind hier nicht abgebildet.", x, 81*mm, width, 11, 14)
    y = pdf_paragraph(c, "Erprobt wurde nur ein 5,99 m langer Transporter von Osten mit Einweiser und zwei Ansätzen. Für den 7,20 m langen Getränkelastwagen liegt keine Schleppkurve vor. Die Skizze ist weder Katasterauszug noch Darstellung von Grundstücksgrenzen.", x, y, width, 11, 14)
    pdf_paragraph(c, "Bezug: Niederschrift vom 09.09.2026. Gemessene Torhöhe 3,10 m; Getränkefahrzeug laut Betreiberangabe 3,35 m hoch.", x, y, width, 10, 13)
    c.setFont("Times-Roman", 8)
    c.drawString(x,15*mm,"LB-18-M1 | Straßenbetrieb | Seite 1")
    c.save()


def update_metadata():
    readme = CASE / "README.md"
    start, end = "<!-- BEGIN native-inventory -->", "<!-- END native-inventory -->"
    table = [start, "| Datei | Inhalt |", "| --- | --- |"]
    table.extend(f"| [{name}]({name}) | {description} |" for name, description in sorted(INVENTORY.items()))
    extras = []
    for extra in sorted(CASE.iterdir()):
        if extra.name[:3] in {"28_", "29_", "30_"} and extra.suffix in {".png", ".jpg", ".xlsx"}:
            extras.append(extra)
            description = {
                "28_": "Nachgebildete interne Schlüsselablage, Stand 22.09.2026 um 15:42 Uhr",
                "29_": "Nachgebildete Chatansicht der vier Nachrichten aus TXT 20; Export 21.09.2026 um 18:06 Uhr",
                "30_": "Arbeitsmappe aus Fahrtenbuch, Schlüsseljournal und Beobachtungen der CSV-Quellen 23 bis 25",
            }[extra.name[:3]]
            table.append(f"| [{extra.name}]({extra.name}) | {description} |")
    table.append(end)
    content = readme.read_text(encoding="utf-8")
    if start in content:
        before, rest = content.split(start, 1)
        _, after = rest.split(end, 1)
        content = before + "\n".join(table) + after
    else:
        anchor = "CSV: UTF-8, Semikolon als Trennzeichen"
        content = content.replace(anchor, "\n".join(table) + "\n\n" + anchor)
    count_start, count_end = "<!-- BEGIN native-count -->", "<!-- END native-count -->"
    native_count = f"Die Akte enthält {27 + len(extras)} native Quelldateien: 11 DOCX, 8 EML, 3 TXT, 3 CSV und 2 PDF"
    if extras:
        images = sum(p.suffix in {".png", ".jpg"} for p in extras)
        sheets = sum(p.suffix == ".xlsx" for p in extras)
        native_count += f", ergänzt um {images} Bildanlagen"
        if sheets:
            native_count += f" und {sheets} XLSX"
    count_block = count_start + "\n" + native_count + ".\n" + count_end
    if count_start in content:
        before, rest = content.split(count_start, 1)
        _, after = rest.split(count_end, 1)
        content = before + count_block + after
    else:
        content = content.replace(start, count_block + "\n\n" + start, 1)
    content = normalize_decimal_headings(content)
    readme.write_text(content, encoding="utf-8")
    (CASE / "rubric.yaml").write_text("""# Ausschließlich technische Existenz- und Mengenprüfungen.
name: Straßenbenutzung Lindenhof Münster
plugin: strassennutzung-genehmigungen
checks:
  - id: r01-readme
    check_type: file_exists
    description: Aktenbezogene README vorhanden
    path: README.md
  - id: r02-native-quellen
    check_type: working_file_count
    description: Mindestens 30 native Quelldateien vorhanden
    min: 30
  - id: r03-fahrten-csv
    check_type: file_exists
    description: Fahrtenexport vorhanden
    path: 23_fahrtenbuch_2026-08-24_bis_09-08.csv
  - id: r04-schluessel-csv
    check_type: file_exists
    description: Schlüsseljournal vorhanden
    path: 24_schluesseljournal_2026-08-21_bis_09-22.csv
  - id: r05-beobachtungen-csv
    check_type: file_exists
    description: Beobachtungsexport vorhanden
    path: 25_beobachtungen_lindenbogen_2026-08-24_bis_09-18.csv
  - id: r06-plan-pdf
    check_type: file_exists
    description: Native Ortsskizze vorhanden
    path: 27_ortsskizze_lindenbogen_2026-09-09.pdf
  - id: r07-gesamt-pdf
    check_type: file_exists
    description: Gesamt-PDF nach dem gesonderten Exportlauf vorhanden
    path: gesamt-pdf/strassennutzung-poller-lieferzufahrt-lindenhof-muenster_gesamt.pdf
  - id: r08-schluessel-bild
    check_type: file_exists
    description: Bildansicht der internen Schlüsselablage vorhanden
    path: 28_schluesselablage_lb17_2026-09-22.png
  - id: r09-chat-bild
    check_type: file_exists
    description: Bildansicht des Chat-Auszuges vorhanden
    path: 29_chat_fruehdienst_2026-09-04.png
  - id: r10-native-arbeitsmappe
    check_type: file_exists
    description: Native Arbeitsmappe vorhanden
    path: 30_betriebsaufzeichnungen_lindenhof.xlsx
""", encoding="utf-8")


def validate():
    assert len(INVENTORY) == 27, len(INVENTORY)
    counts = {suffix: sum(Path(n).suffix == suffix for n in INVENTORY) for suffix in [".docx", ".eml", ".txt", ".csv", ".pdf"]}
    assert counts == {".docx": 11, ".eml": 8, ".txt": 3, ".csv": 3, ".pdf": 2}, counts
    for name in INVENTORY:
        path = CASE / name
        assert path.stat().st_size > 100
        if path.suffix == ".docx":
            doc = Document(path)
            for section in doc.sections:
                assert abs(section.page_width.mm - 210) < .1
                assert abs(section.page_height.mm - 297) < .1
            text = "\n".join(p.text for p in doc.paragraphs)
            assert len(text.split()) >= 230, (name, len(text.split()))
            with ZipFile(path) as archive:
                assert not any("comments" in entry for entry in archive.namelist())
        elif path.suffix == ".pdf":
            reader = PdfReader(path)
            assert len(reader.pages) == 1
            text = "\n".join(p.extract_text() for p in reader.pages)
        elif path.suffix == ".eml":
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            assert all(message[h] for h in ["Date", "From", "To", "Subject", "Message-ID"])
            assert not message.defects
            text = message.get_body(preferencelist=("plain",)).get_content()
            if name.startswith("17_"):
                attachment = list(message.iter_attachments())[0]
                assert attachment.get_payload(decode=True) == (CASE / "26_rechnung_frischefahrt_2026-09-14.pdf").read_bytes()
        else:
            text = path.read_text(encoding="utf-8")
        assert chr(167) not in text, name
        assert any(char in text for char in "äöüÄÖÜß"), name
        for banned in ["Musterlösung", "Lösungsskizze", "Arbeitsauftrag", "Erwartungshorizont", "Diese Testakte wurde"]:
            assert banned not in text, (name, banned)
    assert len(TRIPS) == 24
    total = sum(Decimal(row[8].replace(",", ".")) for row in TRIPS)
    assert total == Decimal("192.00"), total
    assert sum(row[8] == "18,00" for row in TRIPS) == 8
    assert sum(row[8] == "12,00" for row in TRIPS) == 4
    for filename, expected in [("23_fahrtenbuch_2026-08-24_bis_09-08.csv", 24), ("24_schluesseljournal_2026-08-21_bis_09-22.csv", 17), ("25_beobachtungen_lindenbogen_2026-08-24_bis_09-18.csv", 18)]:
        with (CASE / filename).open(encoding="utf-8", newline="") as handle:
            rows = list(csv.DictReader(handle, delimiter=";"))
        assert len(rows) == expected
        assert all(None not in row for row in rows)
    print(f"Lindenhof: {len(INVENTORY)} native Quellen; Formate {counts}; 24 Fahrten, Zusatzentgelte netto {total:.2f} EUR.")
    print("27 Quelldateien erstellt und strukturell geprüft. Layoutprüfung erfolgt separat mit dem DOCX-Renderer.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--csv-only", action="store_true", help="Nur die drei fest definierten CSV-Quellen aufbauen")
    args = parser.parse_args()
    CASE.mkdir(parents=True, exist_ok=True)
    build_csvs()
    if args.csv_only:
        print("Drei finale CSV-Quellen geschrieben: 24 Fahrten, 17 Schlüsselvorgänge, 18 Beobachtungen.")
        return
    build_invoice()
    build_plan()
    build_letters()
    build_emails()
    build_texts()
    update_metadata()
    validate()


if __name__ == "__main__":
    main()
