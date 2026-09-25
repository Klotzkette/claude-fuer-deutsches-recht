#!/usr/bin/env python3
"""Zwei kaufmännische Bauakten. Autor: Klotzkette. Keine zentralen Schreibzugriffe."""
import argparse
import csv
import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime, parseaddr
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from akten_build_runtime import node_binary, serif_font_path
from akten_docx_format import separate_section_headings
from testakte_office_pdf import office_binary
from readme_decimal_headings import normalize_decimal_headings

ROOT = Path(__file__).resolve().parents[1]
ASSETS = Path(os.environ.get('BAUWIRTSCHAFT_ASSETS', '/tmp/bauwirtschaft-assets'))
QA = ASSETS / 'qa-kaufmaennisch'
WAR = 'bauwirtschaft-baumanagement-werkhalle-warendorf'
BAD = 'bauwirtschaft-buchhaltung-bauunternehmen-bad-salzuflen'
NOTICE = ('Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\n'
          'This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.')
STAMP = datetime(2026, 9, 25, 16)
SOURCES = {WAR: [], BAD: []}
def register_fonts():
    pdfmetrics.registerFont(TTFont('TNR', str(serif_font_path())))
    pdfmetrics.registerFont(TTFont('TNRB', str(serif_font_path(bold=True))))


def path(case, name):
    if name not in SOURCES[case]:
        SOURCES[case].append(name)
    p = ROOT / 'testakten' / case / name
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def euro(value):
    return f'{value:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def csvfile(case, name, headers, rows):
    with path(case, name).open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle, delimiter=';')
        writer.writerow(headers)
        writer.writerows(rows)


def txt(case, name, content):
    path(case, name).write_text(content.strip() + '\n', encoding='utf-8')


def doc(case, name, issuer, receiver, title, date, body, table=None):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(1.8), Cm(2.2), Cm(2.2)
    for key in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
        st = d.styles[key]
        st.font.name, st.font.size = 'Times New Roman', Pt(11)
        st.font.color.rgb = RGBColor(0, 0, 0)
        fonts = st.element.get_or_add_rPr().rFonts
        for attribute in list(fonts.attrib):
            if attribute.endswith('Theme'):
                del fonts.attrib[attribute]
        for script in ('ascii','hAnsi','eastAsia','cs'):
            fonts.set(qn('w:'+script),'Times New Roman')
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.line_spacing = 1.05
    d.styles['Title'].font.size = Pt(16)
    d.styles['Title'].font.bold = True
    d.styles['Heading 1'].font.bold = True
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    d.add_paragraph(issuer)
    d.add_paragraph(receiver)
    d.add_paragraph(date)
    d.add_paragraph(title, 'Title')
    for block in body.strip().split('\n\n'):
        heading = block[0].isdigit() and '\n' not in block and len(block) < 75 and not block[1:2].isdigit()
        p = d.add_paragraph(block, 'Heading 1' if heading else 'Normal')
        p.paragraph_format.widow_control = True
    if table:
        t = d.add_table(rows=0, cols=len(table[0]))
        t.style = 'Table Grid'
        for i, row in enumerate(table):
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
                if i == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E8EBED')
                    cell._tc.get_or_add_tcPr().append(shade)
                    for r in cell.paragraphs[0].runs:
                        r.bold = True
    d.core_properties.author = d.core_properties.last_modified_by = 'Klotzkette'
    d.core_properties.title, d.core_properties.language = title, 'de-DE'
    d.core_properties.created = d.core_properties.modified = STAMP
    separate_section_headings(d)
    d.save(path(case, name))


def pdf(case, name, issuer, receiver, title, date, body, table=None):
    target = path(case, name)
    normal = ParagraphStyle('body', fontName='TNR', fontSize=11, leading=14, spaceAfter=8)
    small = ParagraphStyle('small', parent=normal, fontSize=9, leading=11, spaceAfter=3)
    heading = ParagraphStyle('title', parent=normal, fontName='TNRB', fontSize=16, leading=19, spaceAfter=13)
    flow = [Paragraph(escape(issuer).replace('\n', '<br/>'), normal), Spacer(1, 9),
            Paragraph(escape(receiver).replace('\n', '<br/>'), normal), Paragraph(escape(date), normal),
            Paragraph(escape(title), heading)]
    for block in body.strip().split('\n\n'):
        flow.append(Paragraph(escape(block).replace('\n', '<br/>'), normal))
    if table:
        width = A4[0] - 112
        t = Table([[Paragraph(escape(str(v)), small) for v in row] for row in table],
                  colWidths=[width / len(table[0])] * len(table[0]), repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E8EBED')),
                               ('BOX', (0, 0), (-1, -1), .3, colors.HexColor('#D9D9D9')),
                               ('INNERGRID', (0, 0), (-1, -1), .3, colors.HexColor('#D9D9D9')),
                               ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                               ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
        flow += [Spacer(1, 7), t]
    def footer(canvas, document):
        canvas.setFont('TNR', 8)
        canvas.drawString(56, 32, title + ' | ' + str(document.page))
    SimpleDocTemplate(str(target), pagesize=A4, leftMargin=56, rightMargin=56, topMargin=46,
                      bottomMargin=48, title=title, author='Klotzkette').build(flow, onFirstPage=footer, onLaterPages=footer)


def mail(case, name, sender, recipient, date, subject, body, attachment=None, reply=None):
    sender, recipient = sender.replace('.example', '.de'), recipient.replace('.example', '.de')
    message = EmailMessage(policy=SMTP)
    stamp = datetime.fromisoformat(date + '+02:00')
    address = parseaddr(sender)[1]
    message['Return-Path'] = '<' + address + '>'
    destination_domain = parseaddr(recipient)[1].split('@')[1]
    message['Received'] = f'from mail.{address.split("@")[1]} by mx.{destination_domain} with ESMTPS id {name[:2]}0926; {format_datetime(stamp)}'
    message['From'], message['To'] = sender, recipient
    message['Date'], message['Subject'] = format_datetime(stamp), subject
    message['Message-ID'] = f'<{name[:-4]}@{address.split("@")[1]}>'
    message['Content-Language'] = 'de-DE'
    if name == '12_Rechnung_erneut.eml':
        message['X-Archive-Processing'] = 'PDF attachment stamped on receipt; accounting record ER-260908; original invoice retained separately'
    if reply:
        message['In-Reply-To'] = message['References'] = reply
    message.set_content(body.strip() + '\n', charset='utf-8')
    if attachment:
        p = ROOT / 'testakten' / case / attachment
        message.add_attachment(p.read_bytes(), maintype='application', subtype='pdf', filename=p.name)
    path(case, name).write_bytes(message.as_bytes())


W_OWNER = 'Hagedorn Präzisionsteile GmbH\nWerkstraße 18, 48231 Warendorf'
W_CTRL = 'Jana Feldkamp | Projektcontrolling\nHagedorn Präzisionsteile GmbH, Warendorf'
W_BUILDER = 'Westkamp Hallenbau GmbH\nIndustrieweg 7, 59269 Beckum'
B_OWNER = 'Bauunternehmen Mertens GmbH\nSiemensstraße 22, 32105 Bad Salzuflen'
B_SUB = 'Röding Betonbau GmbH\nGewerbering 14, 32657 Lemgo | Steuernummer 329/5812/4071'
B_MAT = 'Lippe Baustoffhandel GmbH\nLagesche Straße 61, 32657 Lemgo | Steuernummer 329/5743/0912'
B_RENT = 'Weser Miettechnik GmbH\nDieselstraße 9, 32052 Herford | Steuernummer 324/5791/2164'
B_ENG = 'Büro Seidel Tragwerksplanung GmbH\nLange Straße 43, 32105 Bad Salzuflen | Steuernummer 313/5701/6410'

TRADES = [
    ['W01', 'Planung', 120000, 118000, 95000, 0],
    ['W02', 'Erdarbeiten', 180000, 190000, 180000, 0],
    ['W03', 'Tragwerk', 640000, 630000, 480000, 0],
    ['W04', 'Hülle', 420000, 405000, 200000, 15000],
    ['W05', 'Technik', 380000, 360000, 140000, 30000],
    ['W06', 'Außenanlagen', 180000, 150000, 0, 35000],
    ['W07', 'Ausstattung', 140000, 100000, 30000, 25000],
    ['W08', 'Baustelleneinrichtung', 140000, 120000, 75000, 28000],
]
# Ein Istbeleg ist ein einzelner gebuchter Kostenbeleg, keine zusätzliche Zahlung.
W_JOURNAL = []
for i, trade in enumerate(TRADES):
    if trade[4]:
        a = {'W05':54000, 'W08':42000}.get(trade[0], round(trade[4] * .4, 2))
        for j, net in enumerate([a, trade[4] - a]):
            W_JOURNAL.append([f'WH-{i+1:02d}-{j+1}', trade[0], f'2026-0{7+j}-28', net, round(net*.19, 2)])
W_PAYMENTS = [260000, 350000, 330000, 185000, 75000]
B_INVOICES = [
    ['ER-260901', 'RB-260901', 'K100', 'BS26-01', 24000, 0, 0, 22800, 1200, 'freigegeben', '2026-09-14'],
    ['ER-260908', 'LB-260908', 'K200', 'BS26-01', 12000, .19, 261.8, 12828.2, 0, 'freigegeben', '2026-09-18'],
    ['GS-260912', 'LB-G260912', 'K200', 'BS26-01', -1000, .19, 0, 0, 0, 'verrechnet', '2026-09-18'],
    ['ER-260915', 'WM-260915', 'K300', 'geteilt', 2500, .19, 0, 2975, 0, 'freigegeben', '2026-09-22'],
    ['ER-260922', 'ST-260922', 'K400', 'BS26-02', 1800, .19, 0, 0, 0, 'ausstehend', '2026-10-06'],
]
B_BANK = [
    ['BK-0901', '2026-09-01', 'Kunde Alte Mühle', 'AR-260801', 60000, 0],
    ['BK-0914', '2026-09-14', 'Röding Betonbau', 'RB-260901', 0, 22800],
    ['BK-0918', '2026-09-18', 'Lippe Baustoffhandel', 'LB-260908 / LB-G260912', 0, 12828.2],
    ['BK-0922', '2026-09-22', 'Weser Miettechnik', 'WM-260915', 0, 2975],
    ['BK-0923', '2026-09-23', 'Lohn August', 'SAM-0826', 0, 15700],
    ['BK-0924A', '2026-09-24', 'Sozialkassen August', 'SV-0826', 0, 7700],
    ['BK-0924B', '2026-09-24', 'Finanzamt', 'LSt-0826', 0, 2000],
    ['BK-0924C', '2026-09-24', 'Weser Miettechnik', 'Mietpark 09', 0, 1785],
]


def warendorf():
    c = WAR
    doc(c, '01_Projektauftrag_Werkhalle.docx', W_OWNER, 'Jana Feldkamp\nProjektcontrolling',
        'Werkhalle Warendorf', 'Warendorf, 25. September 2026', '''1 Anlass und Umfang

Unsere neue Werkhalle mit 1.800 Quadratmetern soll zwei vorhandene CNC-Linien und die Warenausgabe aufnehmen. Wir sind ein inhabergeführter Metallverarbeiter mit 68 Beschäftigten. Bauherrin ist allein die Hagedorn Präzisionsteile GmbH; es gibt keine öffentliche Förderung und keine Vergabe durch eine öffentliche Stelle. Das Grundstück gehört uns bereits und liegt außerhalb des Bauinvestitionsbudgets.

2 Stand der Steuerung

Der Beschluss vom 12. Januar erlaubt 2.300.000,00 EUR netto einschließlich 100.000,00 EUR nicht vergebener Reserve. Die Kostenblätter schließen am 25. September um 16:00 Uhr. Istkosten sind gebuchte Nettoleistungen, nicht Bankabgänge. Verbindliche Bestellungen und zusätzliche, noch unbestellte Restleistungen werden getrennt geführt. Maschinenkauf, Finanzierungskosten und laufende Produktion sind nicht Bestandteil dieser Kostenstelle.

3 Anstehende Entscheidung

Die Trafostation kommt nach heutiger Mitteilung erst am 9. November. Der bisher vorgesehene Produktionsanlauf am 2. November ist damit nicht gesichert. Westkamp bietet eine geänderte Baufolge an; Elektro Breden hat eine zeitlich begrenzte Versorgung angeboten. Die Leistung dieser Versorgung reicht nicht für beide CNC-Linien gleichzeitig. Die Entscheidung liegt bei mir. Bis zur Besprechung am 28. September um 10:00 Uhr ist keine Variante freigegeben.

Bitte führen Sie die laufende Kosten- und Terminübersicht für die Geschäftsführung fort. Eine Erhöhung des Gesamtbudgets oder die Anerkennung des Kranangebots ist damit nicht verbunden. Verhandlungen über Ansprüche gegenüber Lieferanten führt ausschließlich die Geschäftsführung nach gesonderter Abstimmung.

Maren Hagedorn\nGeschäftsführerin''')
    pdf(c, '02_Budgetbeschluss.pdf', W_OWNER, 'Projekt WH26 | Beschlussablage', 'Investitionsfreigabe Werkhalle', '12. Januar 2026',
        '''Die Geschäftsführung genehmigt die nachfolgende Kostenplanung für den Neubau. Alle Werte verstehen sich netto; die Gesellschaft nutzt die Halle ausschließlich für ihre umsatzsteuerpflichtige Fertigung. Grundstück und Maschinen bleiben im bestehenden Anlagenplan.

Die Reserve ist kein Auftrag und keine bereits erwartete Ausgabe. Einzelbeauftragungen bis 25.000,00 EUR innerhalb eines freigegebenen Gewerks zeichnet die Projektleitung gemeinsam mit dem Einkauf. Zusatzleistungen außerhalb der Vergabesumme sowie jede Entnahme aus der Reserve bedürfen der Zustimmung der Geschäftsführerin. Die Mittelbereitstellung erfolgt aus Eigenmitteln und dem Investitionskredit gemäß gesondertem Finanzierungsplan.

Gezeichnet: Maren Hagedorn, Geschäftsführerin. Jana Feldkamp hat den Beschluss am 13. Januar in das Projektregister übernommen.''',
        [['Kostenbereich', 'EUR netto']] + [[r[1], euro(r[2])] for r in TRADES] + [['Reserve', '100.000,00'], ['Gesamtfreigabe', '2.300.000,00']])
    doc(c, '03_Hallenbauvertrag.docx', W_BUILDER, W_OWNER, 'Vertrag über Tragwerk und Gebäudehülle', 'Beckum und Warendorf, 19. Februar 2026', '''1 Gegenstand

Westkamp Hallenbau GmbH errichtet auf dem Grundstück Werkstraße 18 das Stahltragwerk, Dach und Außenwände der 1.800 Quadratmeter großen Werkhalle gemäß Plan WH-04 vom 10. Februar 2026. Zum Tragwerk gehören Fundamente, Stützen, Binder und Aussteifung. Die Gebäudehülle umfasst gedämmte Dach- und Wandpaneele, vier Sektionaltore, Lichtbänder und die Dachentwässerung bis zum Übergabepunkt. Technische Gebäudeausrüstung, Stromversorgung, Außenanlagen und Maschinenmontage sind nicht enthalten.

2 Vergütung und Zahlungen

Die Vergütung beträgt für das Tragwerk 630.000,00 EUR netto und für die Hülle 405.000,00 EUR netto. Umsatzsteuer wird mit 19 Prozent gesondert ausgewiesen. Abschlagsrechnungen weisen den bis zum Stichtag ausgeführten Leistungswert sowie sämtliche vorausgegangenen Rechnungen und Zahlungen aus. Die Bauleitung bestätigt die Mengen. Zahlungen sind 14 Kalendertage nach Eingang der prüfbaren Rechnung zu leisten. Es ist weder Skonto noch ein laufender Sicherheitseinbehalt vereinbart.

3 Ablauf und Schnittstellen

Die Montage beginnt am 11. Mai. Die Halle soll am 16. Oktober wetterdicht und für den Innenausbau zugänglich sein. Die Bauherrin stellt ab dem 12. Oktober die vereinbarte elektrische Anschlussleistung bereit. Die für Maschinenversuche erforderliche Leistung von 450 kVA ist keine Leistung von Westkamp. Behinderungen und mögliche Umstellungen sind am selben Arbeitstag der Bauleitung mitzuteilen; daraus folgt noch keine zusätzliche Vergütungsfreigabe.

4 Änderungen

Zusätzliche Leistungen werden vor Ausführung mit Beschreibung, Preis und Auswirkung auf die Termine angeboten. Die Bauherrin erteilt den Auftrag schriftlich durch ihre Geschäftsführerin. Terminprotokolle ersetzen keine Preisvereinbarung. Teilflächen dürfen nach Abstimmung vorab genutzt werden; hierdurch wird weder eine Schlussabnahme erklärt noch eine fehlende Prüfung ersetzt.

Für Westkamp: Henrik Westkamp. Für Hagedorn: Maren Hagedorn. Beide Parteien haben die Vertragsfassung am 19. Februar 2026 unterzeichnet.''')
    csvfile(c, '04_Kostenrahmen.csv', ['Projekt', 'Gewerk', 'Kostenbereich', 'Budget_netto_EUR'],
            [['WH26', r[0], r[1], euro(r[2])] for r in TRADES] + [['WH26', 'RES', 'Unvergebene Reserve', '100.000,00']])
    pdf(c, '05_Bestellung_Trafostation.pdf', W_OWNER, 'Elektro Breden GmbH\nHülsener Straße 31, 59302 Oelde', 'Bestellung EB 260318', '18. März 2026', '''Wir bestellen die betriebsfertige Trafostation 630 kVA einschließlich Aufstellung, Kabelanschluss und Schutzprüfung für 180.000,00 EUR netto zuzüglich 19 Prozent Umsatzsteuer. Die Station ist Bestandteil der Vergabesumme Technik von 360.000,00 EUR netto; sie darf dort nicht ein zweites Mal angesetzt werden.

Die Lieferung ist für den 5. Oktober, die Inbetriebnahme für den 12. Oktober 2026 vereinbart. Die Abschlagsstufen betragen 54.000,00 EUR netto nach bestätigter Fertigungsfreigabe, 90.000,00 EUR nach Lieferung und 36.000,00 EUR nach dokumentierter Inbetriebnahme. Bis 25. September wurden 54.000,00 EUR fakturiert und bezahlt. Dieser Betrag ist im Ist des Gewerks W05 enthalten.

Die Bestellung umfasst keine mobile Ersatzversorgung. Eine solche Versorgung wäre gesondert anzubieten und durch Maren Hagedorn zu beauftragen. Die Zufahrt für den Tieflader muss spätestens zwei Arbeitstage vor Anlieferung frei sein.

Auftragsannahme durch Elektro Breden am 21. März: Preis, Liefertermin und Zahlungsstufen bestätigt. Einkauf: Leon Rensing.''')
    pdf(c, '07_Lieferfortschreibung.pdf', 'Elektro Breden GmbH\nHülsener Straße 31, 59302 Oelde', W_OWNER,
        'Lieferfortschreibung zur Station EB 260318', '24. September 2026', '''Der Hersteller hat die Endprüfung der Schaltanlage wegen einer noch fehlenden Schutzbaugruppe auf den 5. November verschoben. Wir reservieren den Tieflader für den 9. November und zwei Montageteams für den 10. bis 13. November. Den bisherigen Liefertermin 5. Oktober können wir nicht halten.

Die Inbetriebnahme ist jetzt für den 16. November vorgesehen, sofern die Netzfreigabe vorliegt. Der Preis der bestellten Anlage und die Abschlagsstufen bleiben nach unserer Kalkulation unverändert. Eine mobile Anlage ist darin nicht enthalten. Zum rechtlichen Umgang mit der Terminverschiebung enthält dieses Schreiben keine Vereinbarung.

Unser Disponent bestätigt den Transport am 30. Oktober erneut. Bis dahin bleibt der genannte Termin eine Herstellerzusage, keine bereits erfolgte Lieferung.

Nils Breden\nGeschäftsführer''')
    mail(c, '06_Trafo_Liefermeldung.eml', 'Nils Breden <nils.breden@elektro-breden.example>', 'Jana Feldkamp <j.feldkamp@hagedorn.example>',
         '2026-09-24T11:12:00', 'WH26: neue Lieferwoche Trafostation', '''Guten Tag Frau Feldkamp,

anbei erhalten Sie unsere schriftliche Lieferfortschreibung. Ich habe mit dem Hersteller gesprochen; einen zweiten verfügbaren Schutzschalter kann er derzeit nicht zusagen. Eine frühere Teilanlieferung würde uns bei der Inbetriebnahme nicht helfen.

Für die mobile Anlage habe ich die Anschlusswerte Ihrer beiden Maschinenlinien verwendet. Bitte beachten Sie die Begrenzung auf eine Linie. Wir können die Station nicht allein aufgrund der Terminrunde bestellen. Ich benötige eine gesonderte schriftliche Freigabe. Die Preisbindung für die mobile Anlage endet am Montag um 14:00 Uhr.

Freundliche Grüße
Nils Breden
Elektro Breden GmbH''', '07_Lieferfortschreibung.pdf')
    doc(c, '08_Baubesprechung_25_09.docx', W_CTRL, 'Teilnehmer: Maren Hagedorn, Henrik Westkamp, Nils Breden, Leon Rensing',
        'Baubesprechung Werkhalle', '25. September 2026, 08:30 bis 09:20 Uhr', '''1 Baufortschritt

Die Stahlmontage ist bis Achse 8 abgeschlossen. Zwei Tore werden am 30. September gesetzt. Die Abdichtung des letzten Lichtbandes ist für den 2. Oktober vorgesehen. Westkamp hält den Termin für die wetterdichte Hülle am 16. Oktober weiterhin für erreichbar. Die Mengen für den Septemberabschlag sind am 23. September gemeinsam aufgenommen worden.

2 Energie und Baufolge

Breden nennt den 9. November für die Anlieferung und den 16. November für die Inbetriebnahme. Der bisherige Plan sieht Energie am 12. Oktober und Maschinenversuche bis 27. Oktober vor. Westkamp kann ohne dauerhafte Energie zuerst die Sozialräume und Außenentwässerung fertigstellen. Für die CNC-Versuche bleibt eine leistungsfähige Versorgung erforderlich.

3 Kran und Nachtrag

Kranservice Münsterland verlangt zwölf zusätzliche Miettage und einen weiteren Umsetzvorgang. Rensing bestätigt vier ausgeführte Tage mit Hubbedarf bis heute. Vier weitere Einsatztage sind erst disponiert. Der Nachweis sieht für vier zusätzliche Tage ausschließlich Bereitstellung vor. Ob diese Standtage zu vergüten sind, ist zwischen Einkauf und Vermieter noch nicht abschließend abgestimmt. Eine Preisfreigabe wurde heute nicht erteilt.

4 Festgehaltener Stand

Die Geschäftsführerin entscheidet nach Eingang der aktualisierten Anschlussbestätigung am 28. September. Eine Vollversorgung beider Linien über das Provisorium wurde von keinem Teilnehmer zugesagt. Das Protokoll dokumentiert den Gesprächsstand und verändert keine Bestellung.

Jana Feldkamp\nVersandt an die Teilnehmer am 25. September, 11:10 Uhr''')
    csvfile(c, '11_Belegjournal.csv', ['Beleg', 'Gewerk', 'Buchungsdatum', 'Netto_EUR', 'USt_EUR'],
            [[*r[:3], euro(r[3]), euro(r[4])] for r in W_JOURNAL])
    csvfile(c, '12_Bestellbuch.csv', ['Bestellung', 'Gewerk', 'Auftragnehmer', 'Bestellt_netto_EUR', 'Bestandsdatum'],
            [[f'WH-B{i+1:02d}', r[0], ['Planwerk','Erdmann Tiefbau','Westkamp','Westkamp','Elektro Breden und Partner','Wegewerk','Regalbau Möller','Baustellenservice'][i], euro(r[3]), '2026-08-31'] for i,r in enumerate(TRADES)])
    doc(c, '13_Restkostenerhebung.docx', W_CTRL, 'Geschäftsführung Hagedorn Präzisionsteile GmbH',
        'Restleistungen nach den bestehenden Bestellungen', '25. September 2026', '''Die folgenden Nettoansätze betreffen Leistungen, die noch nicht bestellt sind. Sie enthalten keine offenstehenden Leistungen aus den bestehenden Aufträgen. Das Bestellbuch erfasst die bereits gebundenen Summen vollständig. Die Restkostenerhebung ergänzt diese Beträge und ersetzt nicht die noch ausstehenden Rechnungen.

Für die Hülle fehlen 15.000,00 EUR für zusätzlich gewünschte Innenverkleidung an den Toren. Technik enthält 30.000,00 EUR für die noch zu vergebende Druckluftverteilung. Außenanlagen benötigen weitere 35.000,00 EUR für die Zufahrt im Norden. Die Ausstattung enthält 25.000,00 EUR für die Werkbänke. Die Baustelleneinrichtung benötigt 28.000,00 EUR für planmäßigen Rückbau, Endreinigung und die bereits kalkulierte verlängerte Bauüberwachung. Zu Planung, Erdarbeiten und Tragwerk werden derzeit keine unbestellten Restleistungen erwartet.

Die Ansätze beruhen auf der Mengenrunde vom 23. September. Sie umfassen zusammen 133.000,00 EUR. Die umstrittene Kranforderung von 9.600,00 EUR und das mobile Stromangebot von 36.000,00 EUR sind darin nicht enthalten. Beide Beträge bleiben bis zur Entscheidung gesondert. Die Reserve von 100.000,00 EUR ist weder verbraucht noch nochmals als Restkosten aufzuschlagen.

Jana Feldkamp\nProjektcontrolling''')
    pdf(c, '14_Kranmietvereinbarung.pdf', 'Kranservice Münsterland GmbH\nAm Hafen 17, 48155 Münster', W_OWNER,
        'Mietvereinbarung KM 260511', '7. Mai 2026', '''Wir stellen vom 11. Mai bis 18. September einen Turmdrehkran mit 45 Metern Ausladung auf der Baustelle WH26 bereit. Der Grundpreis einschließlich Aufbau, Abbau und Erstprüfung beträgt 42.000,00 EUR netto. Die Bedienung durch unterwiesenes Personal der Bauunternehmen ist im Grundpreis nicht enthalten. Die Gesellschaft bleibt Mieterin; Westkamp meldet den Einsatzbedarf.

Eine schriftlich vereinbarte Verlängerung kostet 650,00 EUR netto je zusätzlichem Mietwerktag. Samstage und Sonntage werden nicht berechnet. Ein zusätzlich beauftragtes Umsetzen innerhalb der Baustelle kostet 1.800,00 EUR netto. Unvereinbarte Standzeiten und reine Einsatzmeldungen werden vor Rechnungsstellung mit dem Einkauf abgestimmt.

Der Grundpreis ist vollständig im Bestellbuch W08 erfasst. Eine Verlängerung ist zum Abschluss dieser Vereinbarung noch nicht bestellt. Rückgabe und Abbau sind drei Werktage vorher telefonisch anzumelden und schriftlich zu bestätigen.

Für den Vermieter: Anja Kortmann. Für die Mieterin: Leon Rensing und Jana Feldkamp.''')
    pdf(c, '15_Nachtragsangebot_Kran.pdf', 'Kranservice Münsterland GmbH\nAm Hafen 17, 48155 Münster', W_OWNER,
        'Angebot KM N03 zur Verlängerung', '24. September 2026', '''Wir bieten die Verlängerung um zwölf Mietwerktage vom 21. September bis 6. Oktober sowie ein zusätzliches Umsetzen an. Die Berechnung erfasst auch Tage, an denen der Kran für Sie bereitsteht, aber kein Hub angefordert wird. Die Verlängerung wird erst nach Ihrer schriftlichen Annahme verbindlich.

Unser Disponent hat den ursprünglich geplanten Abbau am 18. September aufgrund der telefonischen Baustellenmeldung zurückgestellt. Wir können den Kran nicht tageweise zwischenvermieten. Bitte teilen Sie uns bis 28. September mit, ob der Abbau nun für den 7. Oktober eingeplant werden soll. Das Angebot ist keine Rechnung und enthält keine Zahlungsaufforderung.

Anja Kortmann\nDisposition und Vermietung''', [['Leistung', 'Menge', 'Einzelpreis EUR', 'Netto EUR'], ['Mietwerktage', '12', '650,00', '7.800,00'], ['Umsetzen', '1', '1.800,00', '1.800,00'], ['Summe netto', '', '', '9.600,00'], ['Umsatzsteuer 19 Prozent', '', '', '1.824,00'], ['Gesamt', '', '', '11.424,00']])
    pdf(c, '16_Kran_Einsatznachweis.pdf', W_BUILDER, 'Einkauf Hagedorn Präzisionsteile GmbH', 'Disposition Kranverlängerung', '25. September 2026', '''Die Bauleitung hat die zusätzlichen Mietwerktage wie folgt disponiert. Nur die bis zum 25. September erfassten Hube sind ausgeführt; die späteren Termine sind Bedarfsmeldungen. Die Unterschrift bestätigt die Disposition, nicht die Preisvereinbarung mit dem Vermieter.

Für den 22. September ist der Umsetzvorgang durch den Polier bestätigt. Die vier Tage ohne Hubbedarf sind weiterhin als Standtage eingeplant. Ob der Lieferant diese im Angebot berechnet, entscheidet nicht die Bauleitung.

Ralf Thiele\nPolier Westkamp''', [['Datum', 'Hubbedarf', 'Stand'], ['21.09.2026','ja','ausgeführt'],['22.09.2026','ja, Umsetzen','ausgeführt'],['23.09.2026','nein','Standtag'],['24.09.2026','ja','ausgeführt'],['25.09.2026','ja','ausgeführt'],['28.09.2026','nein','geplant'],['29.09.2026','ja','geplant'],['30.09.2026','ja','geplant'],['01.10.2026','nein','geplant'],['02.10.2026','ja','geplant'],['05.10.2026','ja','geplant'],['06.10.2026','nein','geplant']])
    mail(c, '17_Einkauf_Kran.eml', 'Leon Rensing <einkauf@hagedorn.example>', 'Anja Kortmann <a.kortmann@kranservice.example>', '2026-09-25T10:40:00', 'KM N03: noch keine Preisfreigabe', '''Guten Tag Frau Kortmann,

ich bestätige den Eingang Ihres Angebots. Unsere Bauleitung nennt acht Tage mit Hubbedarf, davon vier noch bevorstehend. Bei vier weiteren Tagen ist nur die Bereitstellung vermerkt. Bitte erläutern Sie, auf welche Absprache Sie sich für die durchgehende Mietzeit stützen.

Den Umsetzvorgang vom 22. September haben wir im Nachweis gefunden. Ich kann Ihnen heute weder die Gesamtsumme bestätigen noch eine Zahlung zusagen. Die Geschäftsführerin bespricht die weitere Baufolge am Montag. Bis dahin bleibt die Verlängerung in unserem Bestellbuch unfreigegeben.

Mit freundlichen Grüßen
Leon Rensing''', '15_Nachtragsangebot_Kran.pdf')
    pdf(c, '18_Angebot_Mobilversorgung.pdf', 'Elektro Breden GmbH\nHülsener Straße 31, 59302 Oelde', W_OWNER, 'Mobile Stromversorgung MB 0925', '25. September 2026', '''Wir bieten eine mobile Versorgung vom 12. Oktober bis einschließlich 22. November für sechs Wochen an. Die maximale nutzbare Leistung beträgt 250 kVA. Damit kann eine CNC-Linie einschließlich der notwendigen Nebenaggregate erprobt und im reduzierten Betrieb genutzt werden. Ein gleichzeitiger Betrieb beider Linien mit zusammen 450 kVA ist ausgeschlossen.

Das Angebot umfasst Anlieferung, Aufbau, Schutzprüfung, Rückbau und die sechs Wochen Miete. Der kalkulierte Betriebsstoffbedarf ist als Pauschale für den vereinbarten Einschichtbetrieb enthalten. Mehrschichtbetrieb und längere Nutzung wären neu zu kalkulieren. Die endgültige Anschlussprüfung am Einspeisepunkt und die Freigabe des Netzbetreibers stehen noch aus.

Wir halten die Anlage bis 28. September um 14:00 Uhr unverbindlich verfügbar. Beauftragung und 30 Prozent Abschlag sind nach schriftlicher Annahme vorgesehen, der Rest nach Rückbau. Dies ist ein Angebot und keine bereits ausgelöste Zahlung.

Nils Breden''', [['Position','EUR netto'],['Transport und Aufbau','6.000,00'],['Sechs Wochen Miete zu 4.000 EUR','24.000,00'],['Prüfung und Rückbau','3.000,00'],['Betriebsstoffpauschale','3.000,00'],['Summe netto','36.000,00'],['Umsatzsteuer 19 Prozent','6.840,00'],['Gesamt','42.840,00']])
    mail(c, '19_Netzanschluss.eml', 'Sven Büttner <anschluss@emsnetz.example>', 'Nils Breden <nils.breden@elektro-breden.example>', '2026-09-25T09:05:00', 'WH26 mobile Versorgung: Schaltbild fehlt', '''Guten Tag Herr Breden,

uns liegt bislang nur die Leistungsangabe 250 kVA vor. Für die Prüfung benötigen wir das einpolige Schaltbild, den Schutzprüfplan und die Verriegelung gegen Rückspeisung. Ohne diese Unterlagen können wir den Anschluss nicht freigeben.

Für die dauerhafte Station ist der 16. November in der Terminreservierung hinterlegt. Eine abschließende Bestätigung ist erst nach Vorlage der Prüfprotokolle möglich. Bitte übersenden Sie das Schaltbild des Provisoriums bis Montag 09:00 Uhr. Wir melden uns anschließend mit dem Prüfergebnis.

Freundliche Grüße
Sven Büttner
Emsnetz Anschlussservice''')
    mail(c, '20_Geschaeftsfuehrung.eml', 'Maren Hagedorn <m.hagedorn@hagedorn.example>', 'Jana Feldkamp <j.feldkamp@hagedorn.example>', '2026-09-25T12:15:00', 'WH26 Entscheidung am Montag', '''Guten Tag Frau Feldkamp,

eine Linie ab Anfang November wäre für die Liefertermine hilfreich, aber nicht gleichbedeutend mit dem vollen Produktionsanlauf. Ich möchte vor der Entscheidung den Anschlussstatus und die Auswirkungen auf unsere September- bis Dezemberzahlungen sehen. Bitte lassen Sie die unverbindliche Kranforderung nicht stillschweigend zu einer Bestellung werden.

Die Mietverlängerung unserer bisherigen Produktionsfläche beträgt 8.000,00 EUR je angefangenen Monat. Für November können wir bis 30. September verlängern. Diese Betriebsausgabe gehört nicht in das Bauinvestitionsbudget. Ohne Provisorium rechnen wir derzeit mit dem Anlauf am 7. Dezember; ob der Vermieter auch Dezember anbietet, klärt der Einkauf noch.

Ich gebe heute weder das Provisorium noch zusätzliche Kranmiete frei.

Maren Hagedorn''')
    # Restzahlungen des Basisforecasts: Istrest 160.000 plus 1.006.000 neue Leistung netto.
    csvfile(c, '21_Zahlungsplan.csv', ['Monat', 'Alte_Istverbindlichkeit_netto_EUR', 'Neue_Leistung_netto_EUR', 'Zahlung_brutto_EUR', 'Planstand'],
            [['2026-09', '160.000,00','0,00','190.400,00','25.09.2026'], ['2026-10','0,00','260.000,00','309.400,00','25.09.2026'],['2026-11','0,00','430.000,00','511.700,00','25.09.2026'],['2026-12','0,00','316.000,00','376.040,00','25.09.2026']])
    pdf(c, '22_Abschlagsrechnung_September.pdf', W_BUILDER + '\nSteuernummer 304/5872/2109', W_OWNER,
        'Abschlagsrechnung WH 260924', '24. September 2026', '''Für Tragwerk und Gebäudehülle rechnen wir den gemeinsam am 23. September aufgenommenen Leistungsstand ab. Leistungszeitraum dieser Abschlagsstufe: 1. bis 23. September 2026. Die kumulierten Beträge enthalten die bisherigen Abschlagsrechnungen. Eine Schlussrechnung liegt noch nicht vor.

Die bis August berechneten 520.000,00 EUR netto sind vollständig bezahlt. Für die Septemberstufe ergibt sich der nachstehende Zahlbetrag. Er ist am 8. Oktober ohne Abzug auf unser bekanntes Geschäftskonto zu leisten. Die Leistung ist im Septemberjournal in W03 mit 96.000,00 EUR und W04 mit 64.000,00 EUR enthalten; die übrigen Augustbuchungen wurden bei der Leistungsfortschreibung übernommen.

Henrik Westkamp\nGeschäftsführer''', [['Berechnung','Netto EUR','USt EUR','Brutto EUR'],['Leistung kumuliert','680.000,00','129.200,00','809.200,00'],['Bisher berechnet und bezahlt','520.000,00','98.800,00','618.800,00'],['Aktueller Abschlag','160.000,00','30.400,00','190.400,00']])
    pdf(c, '23_Leistungsstand_Halle.pdf', W_BUILDER, W_CTRL, 'Leistungsfeststellung Tragwerk und Hülle', '23. September 2026', '''Die gemeinsam aufgenommenen Mengen ergeben den folgenden kumulierten Wert. Tragwerk ist zu 480.000,00 EUR netto ausgeführt; Gebäudehülle zu 200.000,00 EUR. Der Wertnachweis ist Grundlage der Abschlagsrechnung WH 260924. Er umfasst keine technische Gebäudeausrüstung und keine Kranmietverlängerung.

Die Augustzwischenstände wurden für die Rechnung mit 384.000,00 EUR Tragwerk und 136.000,00 EUR Hülle zugrunde gelegt. Die Septemberzunahme beträgt damit 96.000,00 EUR und 64.000,00 EUR. Die Feststellung von Mengen und Leistungswerten stellt keine Abnahme der gesamten Halle dar. Dichtheitsprüfung, Toreinstellung und Dokumentation stehen noch aus.

Bestätigt für Westkamp: Henrik Westkamp. Bestätigt für die Bauherrin: Jana Feldkamp.''', [['Leistung','Auftrag netto EUR','Erbracht netto EUR'],['Tragwerk','630.000,00','480.000,00'],['Gebäudehülle','405.000,00','200.000,00'],['Summe','1.035.000,00','680.000,00']])
    txt(c, '26_ERP_Projektkonto.txt', '''Hagedorn Präzisionsteile GmbH
ERP Projektkonto WH26 | Export 25.09.2026 16:00
Währung EUR | Werte netto | Kostenbelege bis 25.09.2026
Gebuchte Leistungen: 1.200.000,00
Verbindlich bestellt einschließlich bereits gebuchter Leistungen: 2.073.000,00
Bis 25.09.2026 bezahlte Leistungen netto: 1.040.000,00
Zahlungen einschließlich 19 Prozent Umsatzsteuer: 1.237.600,00
Offene gebuchte Verbindlichkeiten netto: 160.000,00
Offene gebuchte Verbindlichkeiten brutto: 190.400,00
Ausstehende Leistungen in Bestellungen sind keine gebuchten Verbindlichkeiten.
Die aktuelle Septemberrechnung wurde am 25.09. nacherfasst. Die frühere Augustlieferung ist im Bestandsjournal enthalten; die Kostenbelegliste wurde entsprechend datiert.
Sachbearbeitung: Jana Feldkamp
Exportumfang: gesamte Kostenstelle WH26, ohne Grundstück und Maschinen.
Ende des Auszugs.''')
    doc(c, '27_Logistik_Baufolge.docx', W_BUILDER, W_OWNER, 'Baufolge bei späterer Stromversorgung', '25. September 2026', '''1 Verfügbare Arbeitsbereiche

Die Sozialräume können mit dem bestehenden Baustrom fertiggestellt werden. Ab 12. Oktober lassen sich dort Türen, Sanitärobjekte und Oberflächen bearbeiten. Die Außenentwässerung kann parallel außerhalb der späteren Tiefladerzufahrt ausgeführt werden. Die Nordzufahrt muss ab 6. November für die Station frei bleiben; dort lagern zurzeit Paneelpakete.

2 Maschinenbereich

Die Maschinen lassen sich mechanisch aufstellen, solange keine elektrische Inbetriebnahme erfolgt. Der Maschinenbauer benötigt nach bestätigter Energieversorgung 15 Kalendertage für die Versuche; anschließend sind sechs Kalendertage für Einweisung und Produktionsfreigabe eingeplant. Ohne Provisorium ergibt sich aus dem Energieansatz 16. November ein geplanter Anlauf am 7. Dezember. Das ist eine Bauablaufannahme, keine Terminbestätigung des Maschinenbauers.

3 Provisorium

Bei einer ab 12. Oktober freigegebenen mobilen Versorgung könnte zunächst nur Linie 1 erprobt werden. Linie 2 folgt nach dem dauerhaften Netzanschluss. Die Arbeitsbereiche lassen sich absperren, ohne die Fluchtwege zu beeinträchtigen. Elektro Breden muss zuvor die Leitungswege und den sicheren Umschaltvorgang planen. Eine Bestellung der Anlage lässt sich aus diesem Ablaufvorschlag nicht ableiten.

Henrik Westkamp\nProjektleiter''')
    pdf(c, '28_Elektro_Schnittstellen.pdf', 'Elektro Breden GmbH\nProjektleitung WH26', W_CTRL, 'Anschlusswerte der Fertigung', '22. September 2026', '''Die beiden Fertigungslinien benötigen einschließlich Absaugung und Kühleinrichtungen gemeinsam 450 kVA. Linie 1 ist mit 220 kVA einschließlich ihrer Nebenaggregate angesetzt. Für die allgemeinen Hallenverbraucher sind weitere 20 kVA vorgesehen. Eine mobile Versorgung mit 250 kVA kann daher Linie 1 mit den Hallenverbrauchern versorgen, nicht beide Fertigungslinien.

Die Angaben beruhen auf den Geräteblättern des Maschinenbauers vom 17. September. Anlaufströme und Schutzkoordination sind vor Anschluss gesondert zu prüfen. Der mobile Aufbau darf nicht parallel zum Netz betrieben werden. Ein verbindlicher Anschlussplan ist noch nicht gezeichnet.

Nils Breden\nElektroplanung''', [['Verbraucher','Planleistung kVA'],['Linie 1 einschließlich Nebenaggregate','220'],['Linie 2 einschließlich Nebenaggregate','230'],['Allgemeine Hallenverbraucher','20']])
    txt(c, '29_Bautagebuch_24_09.txt', '''Westkamp Hallenbau GmbH
Bautagebuch WH26 | 24.09.2026 | Polier Ralf Thiele
Wetter 07:00 Uhr: bedeckt, 13 Grad; 15:00 Uhr: trocken, 18 Grad.
Personal: sechs Monteure, zwei Dachdecker, ein Polier.
Arbeitszeit 07:00 bis 16:00 Uhr, Pause 12:00 bis 12:30 Uhr.
Die letzten zwei Dachpakete wurden an Achse 7 eingehoben. Der Kran war zwischen 08:15 und 10:40 Uhr sowie 13:20 und 14:05 Uhr im Einsatz.
Die Lieferung der Tordichtungen wurde auf 28.09. bestätigt. Der Nordweg ist durch Paneelpaletten belegt. Für die spätere Trafoanlieferung muss die Lagerfläche geräumt werden.
Keine Personenunfälle. Keine Abnahme. Der Einkauf wurde um 15:10 Uhr auf die noch ungeklärte zusätzliche Kranmiete hingewiesen.
Eintrag abgeschlossen um 16:20 Uhr durch Ralf Thiele.''')
    mail(c, '30_Finanzierung.eml', 'Elena Voß <e.voss@hagedorn.example>', 'Jana Feldkamp <j.feldkamp@hagedorn.example>', '2026-09-25T15:15:00', 'Projektkonto WH26 und Abrufplanung', '''Guten Tag Frau Feldkamp,

zum 25. September sind auf dem Projektkonto 512.400,00 EUR verfügbar. Die kumulierte Ausstattung des Kontos betrug 1.750.000,00 EUR, die Zahlungen 1.237.600,00 EUR. Weitere Kreditabrufe von 400.000,00 EUR im Oktober und 600.000,00 EUR im November sind im bestehenden Rahmen vorgesehen. Zins und Tilgung laufen über das Hauptkonto und gehören nicht in diese Übersicht.

Die Abschlagsrechnung WH 260924 über 190.400,00 EUR ist noch unbezahlt. Der Finanzplan zieht sie vorsorglich bereits im September ab, obwohl sie am 8. Oktober fällig ist. Unsere Planwerte enthalten weder die Kranverlängerung noch das Provisorium. Vorsteuererstattungen werden für den betrachteten Zeitraum nicht als zusätzlicher Zahlungseingang angesetzt.

Für die Novemberzahlung an Breden wurde die Verschiebung des Lieferabschlags bereits berücksichtigt. Bitte unterscheiden Sie den Bestand des Baukontos vom noch freien Kostenbudget.

Freundliche Grüße
Elena Voß
Finanzen''')
    # Septemberbelege der Hallenrechnung sind im vollständigen Kostenjournal separat ablesbar.
    revised = []
    for row in W_JOURNAL:
        if row[1] in ('W03', 'W04'):
            continue
        revised.append(row)
    revised += [['WH-03-1','W03','2026-08-28',384000,72960], ['WH-03-2','W03','2026-09-25',96000,18240],
                ['WH-04-1','W04','2026-08-28',136000,25840], ['WH-04-2','W04','2026-09-25',64000,12160]]
    W_JOURNAL[:] = sorted(revised)
    csvfile(c, '11_Belegjournal.csv', ['Beleg','Gewerk','Buchungsdatum','Netto_EUR','USt_EUR'], [[*r[:3],euro(r[3]),euro(r[4])] for r in W_JOURNAL])


def bad_salzuflen():
    c = BAD
    doc(c, '01_Uebergabe_Buchhaltung.docx', B_OWNER, 'Nora Brinkmann\nKaufmännische Leitung', 'Belegprüfung September', '25. September 2026, 16:00 Uhr', '''1 Gegenstand

Wir übergeben die fünf Buchungsbelege des ausgewählten Septemberstapels mit den Leistungsnachweisen und Zahlungen. Der Stapel betrifft den Anbau einer Werkstatt in Bad Salzuflen und die Sanierung einer Ladenfläche in Lage. Die Unterlagen sind zur laufenden Rechnungsprüfung bestimmt. Sie sind weder ein vollständiger Jahresabschluss noch Grundlage für eine umfassende Wirtschaftsprüfung. Ausgangsrechnungen, Anlagevermögen und frühere Monate sind nur insoweit enthalten, wie sie einen Bankposten dieses Auszugs erklären.

2 Abgrenzung

Die Lohnzusammenfassung betrifft August und wurde im September bezahlt. Ihre Projektstunden dürfen den Septemberrechnungen zur Kostenfortschreibung zugerechnet werden, sind aber keine Septemberlohnabrechnung. Der Bankauszug enthält sämtliche Umsätze vom 1. bis 25. September auf dem Betriebskonto. Der Zahlungseingang von 60.000,00 EUR betrifft eine bereits im August gebuchte Kundenrechnung; dieser Debitor ist nicht Teil des Kreditorenstapels.

3 Aktueller Arbeitsstand

Die Rechnungsmappe enthält auch den erneut übersandten Baustoffbeleg. Unser Eingangsregister führt den Lieferantenbeleg unter ER-260908. Bei Seidel fehlt die Leistungsbestätigung des Bauleiters; die Rechnung wurde dennoch als Verbindlichkeit erfasst. Für den Bankposten mit Verwendungszweck Mietpark 09 hat das Büro noch keine Rechnung zugeordnet. Nora Brinkmann hat deshalb keine zweite Mietzahlung im Lieferantenkonto ausgeglichen.

Bitte besprechen Sie den beleggestützten Zahlungs- und Projektstand mit mir. Bankaufträge werden weiterhin ausschließlich im Vieraugenverfahren durch Geschäftsführung und kaufmännische Leitung freigegeben.

Tobias Mertens\nGeschäftsführer''')
    csvfile(c, '02_Projektstamm.csv', ['Projekt','Bezeichnung','Bauort','Bauleitung','Kostenbudget_netto_EUR'],
            [['BS26-01','Werkstattanbau Fricke','Bad Salzuflen','Lea Tönnies','180.000,00'],['BS26-02','Ladensanierung Westtor','Lage','Jan Hellwig','95.000,00'],['BS25-08','Alte Mühle abgeschlossen','Lemgo','Lea Tönnies','120.000,00']])
    doc(c, '03_Subunternehmervertrag.docx', B_OWNER, B_SUB, 'Auftrag Betonarbeiten BS26 01', '17. August 2026', '''1 Leistungsumfang

Röding Betonbau GmbH übernimmt die Herstellung der bewehrten Bodenplatte am Werkstattanbau Fricke, einschließlich Schalung, Betonieren, Verdichten und Oberflächenbearbeitung. Die abzurechnende Menge beträgt 200 Quadratmeter bei einem Einheitspreis von 120,00 EUR netto. Die Leistung wird vom 24. bis 31. August ausgeführt und nach gemeinsamem Aufmaß abgerechnet. Materiallieferungen ohne Einbau sind kein gesonderter Teil dieses Auftrags.

2 Abrechnung

Die Auftragssumme beträgt 24.000,00 EUR. Der Leistungsempfänger erbringt nachhaltig Bauleistungen; seine Bescheinigung ist dem Auftrag beigefügt. Die Abrechnung erfolgt ohne gesonderten Umsatzsteuerausweis mit dem Hinweis auf die Steuerschuldnerschaft des Leistungsempfängers. Die Freistellungsbescheinigung des Auftragnehmers ist bis zum 31. Dezember 2026 gültig und liegt dem Auftraggeber vor.

3 Zahlung und Sicherheit

Die Zahlung wird 14 Tage nach Rechnungseingang und Leistungsbestätigung fällig. Skonto ist nicht vereinbart. Der Auftraggeber behält als vertragliche Sicherheit fünf Prozent der Nettoabrechnungssumme ein. Der verbleibende Zahlbetrag von 95 Prozent wird auf das im Lieferantenstamm hinterlegte Konto überwiesen. Die Sicherheit bleibt als offene Verbindlichkeit bestehen. Sie wird nach bestätigter mängelfreier Abnahme und Eingang einer vereinbarten Ablösebürgschaft ausgezahlt; der Auftragnehmer hat zum heutigen Stand keine Bürgschaft gestellt.

4 Leistungskontrolle

Die Bauleitung bestätigt Mengen und Leistungszeitraum. Änderungen sind vor Ausführung schriftlich zu vereinbaren. Die Zahlungsfreigabe der kaufmännischen Leitung ersetzt weder die Abnahme noch einen Nachtrag. Zum Vertragsabschluss sind keine zusätzlichen Leistungen beauftragt.

Für Mertens: Tobias Mertens. Für Röding: Sven Röding. Vertragsfassung beiderseits bestätigt am 17. August 2026.''')
    pdf(c, '04_Aufmass_Bodenplatte.pdf', B_SUB, B_OWNER, 'Aufmaß RB 0831', '31. August 2026', '''Die Bodenplatte im Werkstattanbau Fricke wurde vom 24. bis 31. August hergestellt. Aus dem gemessenen Rechteck von 20,00 Metern mal 10,00 Metern ergeben sich 200,00 Quadratmeter. Die Schalung, Bewehrungsarbeiten, Betonage und vereinbarte Oberflächenbearbeitung sind abgeschlossen.

Die Bauleiterin bestätigt die abrechenbare Menge und den Leistungszeitraum. Die gemeinsame Schlussabnahme des Gesamtbauteils ist für den 30. September vorgesehen. Dieser Nachweis enthält keine Freigabe zur Auszahlung des Sicherheitseinbehalts.

Aufgenommen: Sven Röding. Bestätigt: Lea Tönnies am 1. September 2026.''', [['Leistung','Menge','Einheitspreis EUR','Netto EUR'],['Bodenplatte einschließlich Nebenleistungen','200 m²','120,00','24.000,00']])
    pdf(c, '05_Rechnung_RB_260901.pdf', B_SUB, B_OWNER, 'Rechnung RB 260901', '1. September 2026', '''Wir berechnen die Bodenplatte am Projekt BS26-01 gemäß Auftrag vom 17. August und Aufmaß RB 0831. Leistungszeitraum ist der 24. bis 31. August 2026. Die berechnete Leistung wurde vollständig durch unser Unternehmen ausgeführt.

Der Rechnungsbetrag beträgt 24.000,00 EUR. Steuerschuldnerschaft des Leistungsempfängers nach Paragraf 13b Absatz 2 Nummer 4 in Verbindung mit Absatz 5 UStG. Es wird keine Umsatzsteuer an uns gezahlt. Der Betrag ist kein steuerfreier Umsatz.

Nach dem vereinbarten Sicherheitseinbehalt von 1.200,00 EUR verbleiben 22.800,00 EUR zur Zahlung am 14. September. Skonto ist nicht vereinbart. Bitte verwenden Sie die Rechnungsnummer RB-260901. Unsere dem Auftrag beigefügte Freistellungsbescheinigung gilt auch für diese Zahlung.

Sven Röding\nGeschäftsführer''', [['Leistung','Menge','Preis EUR','Betrag EUR'],['Bodenplatte','200 m²','120,00','24.000,00'],['Rechnungssumme','','','24.000,00'],['Sicherheit fünf Prozent','','','1.200,00'],['Überweisungsbetrag','','','22.800,00']])
    pdf(c, '06_Freistellungsbescheinigung.pdf', 'Büroabschrift aus der Lieferantenakte\nRöding Betonbau GmbH', B_OWNER, 'Freistellungsbescheinigung Röding Betonbau', 'Übernommen am 17. August 2026', '''Die im Einkauf hinterlegte Bescheinigung wurde durch das Finanzamt Lemgo am 12. Januar 2026 ausgestellt. Sie bezeichnet die Röding Betonbau GmbH, Gewerbering 14, 32657 Lemgo, Steuernummer 329/5812/4071, als Leistende. Die Geltungsdauer reicht vom 12. Januar bis zum 31. Dezember 2026. Die Freistellung nach Paragraf 48b Absatz 1 EStG ist nicht auf ein einzelnes Bauvorhaben oder einen bestimmten Leistungsempfänger beschränkt.

Die Originalbescheinigung liegt im Lieferantenarchiv unter K100-FSB-2026. Diese Abschrift wurde von Nora Brinkmann anhand der dort abgelegten Bescheinigung angefertigt. Es handelt sich nicht um einen vom Bauunternehmen selbst erlassenen Verwaltungsakt. Bis zur Zahlung am 14. September ist dem Einkauf keine Aufhebung oder Einschränkung zugegangen. Der Zahlungslauf verweist auf diesen Archivbeleg.

Nora Brinkmann\nLieferantenbuchhaltung, 14. September 2026''')
    pdf(c, '07_Bescheinigung_Bauleistungen.pdf', 'Büroabschrift aus der Steuerakte\nBauunternehmen Mertens GmbH', B_SUB, 'Bescheinigung über nachhaltige Bauleistungen', 'Kopie zum Auftrag vom 17. August 2026', '''Die Steuerakte enthält eine Bescheinigung nach dem Vordruck USt 1 TG für die Bauunternehmen Mertens GmbH, Siemensstraße 22, 32105 Bad Salzuflen, Steuernummer 313/5872/1940. Aussteller ist das Finanzamt Lemgo. Die Bescheinigung wurde am 8. Januar 2026 ausgestellt und ist bis zum 31. Dezember 2028 befristet. Sie bestätigt die nachhaltige Erbringung von Bauleistungen im Sinne des Paragrafen 13b Absatz 2 Nummer 4 UStG.

Die Bescheinigung lag bei Ausführung der Betonarbeiten im August vor. Die Kopie wurde Röding am 17. August zusammen mit dem Auftrag übermittelt. In der Steuerakte ist zum 25. September weder eine Rücknahme noch ein Widerruf vermerkt. Das Dokument betrifft die Eigenschaft des Leistungsempfängers und ersetzt nicht das Aufmaß oder den Nachweis der konkreten Bauleistung.

Übertragen aus Archivbeleg ST-U1TG-2026 durch Nora Brinkmann am 17. August 2026.''')
    pdf(c, '08_Rechnung_LB_260908.pdf', B_MAT, B_OWNER, 'Rechnung LB 260908', '8. September 2026', '''Wir berechnen die am 7. September mit Lieferschein LS-0907 an die Baustelle BS26-01 gelieferten 400 Mauersteineinheiten. Der Preis beträgt 30,00 EUR je Einheit. Gegenstand ist ausschließlich die Lieferung; Versetzen, Einbau und sonstige Bauarbeiten werden von uns nicht ausgeführt.

Zahlbar bis 22. September ohne Abzug oder bis 18. September mit zwei Prozent Skonto auf den nach verrechneten Rücknahmen verbleibenden Bruttobetrag. Bitte geben Sie LB-260908 an. Die Übermittlung als PDF ist mit Ihrer Buchhaltung vereinbart. Die Ware bleibt nach unseren vereinbarten Lieferbedingungen bis zur vollständigen Bezahlung unser Eigentum.

Lippe Baustoffhandel GmbH\nDebitorenbuchhaltung''', [['Ware','Menge','Preis netto EUR','Netto EUR'],['Mauersteineinheiten gemäß LS-0907','400','30,00','12.000,00'],['Umsatzsteuer 19 Prozent','','','2.280,00'],['Rechnungsbetrag','','','14.280,00']])
    stamp = io.BytesIO()
    cv = canvas.Canvas(stamp, pagesize=A4)
    cv.setStrokeColor(colors.HexColor('#7B8589'))
    cv.setFillColor(colors.HexColor('#4E5B60'))
    cv.rect(316, 75, 223, 38)
    cv.setFont('TNR', 9)
    cv.drawString(325, 98, 'Posteingang 16.09.2026 10:24 Uhr')
    cv.drawString(325, 84, 'Archiv ER-260908 | erneute Übersendung')
    cv.save()
    copied = PdfReader(path(c, '08_Rechnung_LB_260908.pdf'))
    copied.pages[0].merge_page(PdfReader(io.BytesIO(stamp.getvalue())).pages[0])
    writer = PdfWriter()
    writer.append(copied)
    writer.add_metadata({'/Author':'Klotzkette','/Title':'Rechnung LB 260908 mit Eingangskennung','/Creator':'Klotzkette'})
    writer.write(path(c, '09_Lieferantenmail_Anhang_908.pdf'))
    pdf(c, '10_Lieferschein_LS_0907.pdf', B_MAT, 'Baustelle Fricke | BS26-01\nWerkstattstraße 8, Bad Salzuflen', 'Lieferschein LS 0907', '7. September 2026', '''Geliefert wurden 400 Mauersteineinheiten gemäß Bestellung M-0828. Die Ware wurde um 09:10 Uhr auf der gekennzeichneten Lagerfläche abgesetzt. Einbauleistungen sind nicht Bestandteil der Lieferung. Lea Tönnies bestätigte die Anzahl der Packeinheiten ohne Vorbehalt.

Am 10. September wurden die ungeöffneten Restpakete im vereinbarten Rücknahmewert von 1.000,00 EUR netto abgeholt. Rücknahmeschein RT-0910 ist auf dieser Lieferscheinausfertigung mitgeführt. Die Bewertung wurde im Einkauf am 11. September bestätigt; die gesonderte Rechnungskorrektur folgt durch die Debitorenbuchhaltung.

Anlieferung quittiert: Lea Tönnies. Rücknahme quittiert: Jonas Evers, Fahrer Lippe Baustoffhandel.''')
    pdf(c, '11_Gutschrift_LB_G260912.pdf', B_MAT, B_OWNER, 'Rechnungskorrektur LB G260912', '12. September 2026', '''Für die Rücknahme vom 10. September mindern wir unsere Rechnung LB-260908 um 1.000,00 EUR netto zuzüglich 190,00 EUR Umsatzsteuer. Der Gutschriftsbetrag beträgt 1.190,00 EUR. Es handelt sich um eine Rechnungskorrektur des Lieferanten, nicht um eine Abrechnung durch den Leistungsempfänger.

Bitte verrechnen Sie die Gutschrift mit der offenen Rechnung. Es erfolgt keine separate Auszahlung. Der verbleibende Bruttobetrag beträgt 13.090,00 EUR. Die Skontofrist der Ursprungsrechnung bleibt unverändert; bei Eingang bis 18. September dürfen zwei Prozent von diesem verbleibenden Betrag abgezogen werden.

Lippe Baustoffhandel GmbH\nDebitorenbuchhaltung''')
    mail(c, '12_Rechnung_erneut.eml', 'Lippe Debitoren <debitoren@lippe-baustoff.example>', 'Nora Brinkmann <n.brinkmann@mertens-bau.example>', '2026-09-16T10:24:00', 'LB-260908 erneut im Anhang', '''Guten Tag Frau Brinkmann,

wie telefonisch besprochen senden wir die bereits am 8. September ausgestellte Rechnung unverändert noch einmal. Dies ist keine neue Forderung. Die Gutschrift LB-G260912 über 1.190,00 EUR verrechnen Sie bitte vor dem Skontoabzug.

Bei Zahlungseingang bis 18. September beträgt der Überweisungsbetrag 12.828,20 EUR. Der Skontoabzug von 261,80 EUR setzt sich aus 220,00 EUR Nettominderung und 41,80 EUR Umsatzsteuerkorrektur zusammen. Eine zusätzliche Auszahlung der Gutschrift erfolgt nicht. Als Zahlungsreferenz reicht die ursprüngliche Rechnungsnummer aus.

Freundliche Grüße
Kathrin Hesse
Lippe Baustoffhandel GmbH''', '09_Lieferantenmail_Anhang_908.pdf')
    pdf(c, '13_Rechnung_Miettechnik.pdf', B_RENT, B_OWNER, 'Rechnung WM 260915', '15. September 2026', '''Für die im Zeitraum 1. bis 14. September ohne Bedienpersonal überlassenen Rüttelplatten und Verdichter berechnen wir 2.500,00 EUR netto. Die Geräte wurden durch Ihr Personal bedient. Unser Unternehmen hat keine Arbeiten am Bauwerk erbracht.

Die Rückgabe wurde am 14. September bestätigt. Die Aufteilung auf Ihre Baustellen entnehmen Sie bitte dem zugehörigen Mietnachweis WM-M0914: 1.500,00 EUR entfallen auf BS26-01 und 1.000,00 EUR auf BS26-02. Der Rechnungsbetrag ist am 22. September ohne Skonto fällig. Ein Sicherheitseinbehalt ist nicht vereinbart.

Weser Miettechnik GmbH\nBuchhaltung''', [['Abrechnung','EUR'],['Miete netto','2.500,00'],['Umsatzsteuer 19 Prozent','475,00'],['Rechnungsbetrag','2.975,00']])
    pdf(c, '14_Mietnachweis.pdf', B_RENT, B_OWNER, 'Mietnachweis WM M0914', '14. September 2026', '''Die nachstehenden Geräte wurden ohne Bedienpersonal vermietet. Die Geräte waren bei Rückgabe vollständig und ohne festgestellte neue Beschädigung. Miettage wurden projektbezogen erfasst; Transport ist in den Tagespreisen enthalten.

Dieser Nachweis gehört ausschließlich zur Rechnung WM-260915. Weitere Septembermieten oder Anzahlungen werden hier nicht abgerechnet. Der Empfang wurde durch Lea Tönnies für BS26-01 und Jan Hellwig für BS26-02 bestätigt.

Rücknahme: Enno Weber, Weser Miettechnik.''', [['Projekt','Gerät','Tage','Tagessatz EUR','Netto EUR'],['BS26-01','Rüttelplatte R7','10','150,00','1.500,00'],['BS26-02','Verdichter V3','10','100,00','1.000,00']])
    pdf(c, '15_Rechnung_Tragwerksplanung.pdf', B_ENG, B_OWNER, 'Rechnung ST 260922', '22. September 2026', '''Wir berechnen für die Ladensanierung BS26-02 die überarbeitete statische Berechnung der Türöffnung sowie den Baustellentermin vom 18. September. Leistungszeitraum ist der 11. bis 18. September 2026. Die Berechnung wurde an Jan Hellwig übersandt. Unsere Tätigkeit umfasst ausschließlich Planung und Beratung, keine Ausführung von Bauarbeiten.

Die Rechnung ist am 6. Oktober ohne Abzug fällig. Ein Skonto oder Sicherheitseinbehalt ist nicht vereinbart. Bitte verwenden Sie ST-260922 als Zahlungsreferenz. Für Rückfragen zur Berechnungsfassung steht Ihnen Dipl.-Ing. Paula Seidel zur Verfügung.

Büro Seidel Tragwerksplanung GmbH''', [['Leistung','Netto EUR'],['Überarbeitung Berechnung, pauschal','1.500,00'],['Baustellentermin, pauschal','300,00'],['Summe netto','1.800,00'],['Umsatzsteuer 19 Prozent','342,00'],['Rechnungsbetrag','2.142,00']])
    mail(c, '16_Freigabe_Seidel.eml', 'Nora Brinkmann <n.brinkmann@mertens-bau.example>', 'Jan Hellwig <j.hellwig@mertens-bau.example>', '2026-09-25T08:35:00', 'ST-260922: Leistungsbestätigung ausstehend', '''Guten Morgen Jan,

ich habe ST-260922 mit 2.142,00 EUR als Verbindlichkeit erfasst. Im Zahlungsfeld steht noch keine Freigabe. Bitte bestätige, ob die revidierte Berechnung eingegangen ist und zum tatsächlich ausgeführten Öffnungsmaß passt. Die Rechnung hängt noch einmal an.

Bis zu Deiner Rückmeldung bleibt sie außerhalb des Zahlungslaufs. Die Fälligkeit am 6. Oktober ist eingetragen. Es geht nicht um eine bereits festgestellte Minderleistung; mir fehlt die Bestätigung zum Dokumentstand. Bitte nenne bei Deiner Antwort die Plannummer und das Datum der geprüften Fassung, damit ich sie im Belegarchiv zuordnen kann.

Viele Grüße
Nora''', '15_Rechnung_Tragwerksplanung.pdf')
    pdf(c, '17_Lohnzusammenfassung_August.pdf', 'Lohnbüro Mertens\nAbrechnungskreis gewerbliche Beschäftigte', 'Nora Brinkmann | Buchhaltung', 'Lohnzusammenfassung August', '20. September 2026', '''Die Zusammenfassung betrifft zehn gewerbliche Beschäftigte und 700 abgerechnete Projektstunden des Monats August. Je Beschäftigtem sind 70 Stunden und 2.100,00 EUR Bruttolohn enthalten. Die gesamte Arbeitgeberbelastung von 4.400,00 EUR wurde vom Lohnbüro den Projekten proportional zum Bruttolohn zugeordnet.

Die Projektkosten betragen 25.400,00 EUR. Die Nettoauszahlung von 15.700,00 EUR wurde am 23. September ausgeführt. Die zusammengefassten Zahlungen an Sozialkassen von 7.700,00 EUR und an das Finanzamt von 2.000,00 EUR wurden am 24. September gebucht. Die SV-Zahlung enthält 3.300,00 EUR Arbeitnehmeranteil und 4.400,00 EUR Arbeitgeberbelastung. Weitere Umlagen sind in diesem Abrechnungskreis nicht separat zu zahlen.

Einzelabrechnungen bleiben in der Personalakte. Die mitgelieferten Personalnummern dienen nur der Projektzuordnung. Diese Übersicht enthält keine Aussage über die Septemberlöhne.

Erstellt: Sabine Krüger, Lohnbüro''', [['Position','EUR'],['Bruttolohn','21.000,00'],['Arbeitnehmer SV','3.300,00'],['Lohnsteuer','2.000,00'],['Nettoauszahlung','15.700,00'],['Arbeitgeberbelastung','4.400,00'],['Projektkosten','25.400,00']])
    csvfile(c, '18_Projektstunden_August.csv', ['Personalnummer','Projekt','Monat','Stunden','Stundenlohn_EUR','AG_Belastung_EUR'],
            [[f'P{i+1:03d}', 'BS26-01' if i < 6 else 'BS26-02', '2026-08', '70,00','30,00','440,00'] for i in range(10)])
    csvfile(c, '19_Bankumsatz.csv', ['Bankreferenz','Valuta','Gegenpartei','Verwendungszweck','Haben_EUR','Soll_EUR'],
            [[*r[:4],euro(r[4]),euro(r[5])] for r in B_BANK])
    csvfile(c, '22_Eingangsjournal.csv', ['Buchung','Lieferantenbeleg','Kreditor','Projekt','Netto_EUR','USt_EUR','Belegbetrag_EUR'],
            [[*r[:4],euro(r[4]),euro(round(r[4]*r[5],2)),euro(round(r[4]*(1+r[5]),2))] for r in B_INVOICES])
    csvfile(c, '23_OPOS_25_09.csv', ['Kreditor','Beleg','Belegbetrag_EUR','Zahlung_EUR','Skonto_EUR','Gutschrift_EUR','Offen_EUR'],
            [['K100','RB-260901','24.000,00','22.800,00','0,00','0,00','1.200,00'],['K200','LB-260908','14.280,00','12.828,20','261,80','1.190,00','0,00'],['K300','WM-260915','2.975,00','2.975,00','0,00','0,00','0,00'],['K400','ST-260922','2.142,00','0,00','0,00','0,00','2.142,00']])
    doc(c, '26_Zahlungslauf_Freigabe.docx', B_OWNER, 'Nora Brinkmann und Tobias Mertens', 'Freigaben zum Zahlungslauf', '22. September 2026', '''1 Bereits ausgeführte Zahlungen

Die Zahlung an Röding vom 14. September beträgt 22.800,00 EUR. Die Rechnung über 24.000,00 EUR ist freigegeben; 1.200,00 EUR bleiben als vertragliche Sicherheit offen. Die Freistellungsbescheinigung war bei Zahlung gültig. Eine Ablösebürgschaft liegt nicht vor. Ein Umsatzsteuerbetrag wird dem Lieferanten nicht überwiesen.

Die Baustoffrechnung wurde nach Verrechnung der Gutschrift mit 12.828,20 EUR am 18. September ausgeglichen. Der Lieferant hat die Skontobasis von 13.090,00 EUR bestätigt. Die zweite per E-Mail eingegangene PDF trägt dieselbe Rechnungsnummer und wird unter demselben Eingangsbeleg abgelegt. Es besteht kein zweiter Zahlungsauftrag.

2 Heutiger Lauf

Zur Freigabe steht die Rechnung WM-260915 über 2.975,00 EUR. Der Mietnachweis ist mit beiden Bauleitern abgestimmt. Der Betrag wird ohne Skonto oder Einbehalt gezahlt. Der Lauf enthält keine weiteren Mietzahlungen und keine Rechnung des Büros Seidel.

Nora Brinkmann bestätigt den Belegabgleich. Tobias Mertens bestätigt die Bankfreigabe am 22. September um 11:05 Uhr. Der Bankstatus lautet ausgeführt.''')
    mail(c, '27_Nachfrage_Mietpark.eml', 'Nora Brinkmann <n.brinkmann@mertens-bau.example>', 'Weser Buchhaltung <buchhaltung@weser-miettechnik.example>', '2026-09-25T09:30:00', 'Überweisung 1.785,00 EUR vom 24. September', '''Guten Tag,

auf unserem Bankkonto ist am 24. September eine weitere Überweisung über 1.785,00 EUR an Ihr Unternehmen mit dem Text Mietpark 09 ausgeführt worden. WM-260915 über 2.975,00 EUR wurde bereits am 22. September beglichen. Im vorliegenden Mietnachweis finde ich keinen weiteren Betrag.

Bitte teilen Sie uns mit, welche Rechnung oder Anzahlung Sie dem Betrag zugeordnet haben, und senden Sie den Beleg. Wir führen den Umsatz bis zur Rückmeldung auf dem Verrechnungskonto 1590; ein Projekt ist noch nicht zugeordnet. Eine Rückzahlung fordern wir mit dieser Anfrage noch nicht an.

Mit freundlichen Grüßen
Nora Brinkmann''')
    closing = 85000 + sum(r[4]-r[5] for r in B_BANK)
    pdf(c, '28_Kontoauszug_September.pdf', 'Lippische Gewerbebank eG\nGeschäftskonto Bauunternehmen Mertens GmbH', B_OWNER,
        'Kontoauszug September bis 25 September', '25. September 2026, 16:00 Uhr',
        f'''Konto 471108 | Währung EUR. Anfangssaldo am 1. September: 85.000,00 EUR. Die unten aufgeführten acht Umsätze sind alle Buchungen bis zum Auszugszeitpunkt. Es liegen keine vorgemerkten Umsätze vor. Der verfügbare Guthabenstand entspricht dem Buchsaldo von {euro(closing)} EUR; eine Kreditlinie wird in dieser Ansicht nicht berücksichtigt.

Der Auszug ist elektronisch erstellt und gilt ohne Unterschrift. Wertstellung und Buchungstag stimmen bei den aufgeführten Umsätzen überein.''',
        [['Referenz','Datum','Haben EUR','Soll EUR']] + [[r[0],r[1],euro(r[4]),euro(r[5])] for r in B_BANK] + [['Schlusssaldo','25.09.2026',euro(closing),'']])
    txt(c, '29_Telefonnotiz_Miettechnik.txt', '''Bauunternehmen Mertens GmbH
Telefonnotiz | Nora Brinkmann | 25.09.2026, 13:40 Uhr
Gesprächspartner: Enno Weber, Weser Miettechnik GmbH.
Herr Weber bestätigt den Eingang der Nachricht von heute Vormittag. Er findet WM-260915 in der Debitorenliste mit dem Zahlungseingang vom 22.09. Der zweite Zahlungseingang vom 24.09. ist im Tageskonto sichtbar, aber noch ohne Rechnungsverknüpfung.
Die Disposition prüft, ob eine zusätzliche Reservierung für Lage ausgelöst wurde. Herr Weber kann weder eine Bestellung noch eine Rechnung dazu benennen. Er kündigt eine Rückmeldung für den 28.09. an.
Ich habe keine Verrechnung mit ST-260922 und keine Erstattung vereinbart. Konto 1590 bleibt bis zum Belegzugang unverändert. Projektzuordnung bleibt leer.
Ende der Notiz.''')
    pdf(c, '30_Kundenkonto_Alte_Muehle.pdf', B_OWNER, 'Kundenkonto 12017 | Alte Mühle Gewerberäume GmbH', 'Zahlungsavis zur Augustrechnung', '1. September 2026', '''Der Zahlungseingang von 60.000,00 EUR am 1. September wird der bereits im August gebuchten Abschlagsrechnung AR-260801 zugeordnet. Die Rechnung betrifft das abgeschlossene Ausbauvorhaben Alte Mühle und gehört nicht zu den Projekten BS26-01 und BS26-02.

Vor der Zahlung betrug der offene Betrag dieser Rechnung 60.000,00 EUR. Nach Zuordnung des Zahlungseingangs ist sie vollständig ausgeglichen. Die Einzahlung stellt keinen zusätzlichen Erlös des ausgewählten Septemberstapels dar. Dieser Avis erklärt ausschließlich den Bankumsatz BK-0901; weitere Debitorenbelege sind in der vorliegenden Zusammenstellung nicht enthalten.

Nora Brinkmann\nDebitorenbuchhaltung''')


def metadata_pdf(data, title):
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(data)))
    writer.add_metadata({'/Author':'Klotzkette','/Title':title,'/Creator':'Klotzkette'})
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def office_metadata():
    for case in (WAR, BAD):
        directory = ROOT/'testakten'/case
        for sidecar in directory.glob('*.inspect.ndjson'):
            shutil.move(sidecar, QA / sidecar.name)
        for workbook in directory.glob('*.xlsx'):
            output = io.BytesIO()
            with zipfile.ZipFile(workbook) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
                has_core = 'docProps/core.xml' in source.namelist()
                for entry in source.infolist():
                    data = source.read(entry.filename)
                    if entry.filename == 'docProps/core.xml':
                        root = ET.fromstring(data)
                        for uri, name in [('http://purl.org/dc/elements/1.1/','creator'),('http://schemas.openxmlformats.org/package/2006/metadata/core-properties','lastModifiedBy'),('http://purl.org/dc/elements/1.1/','language')]:
                            el = root.find(f'{{{uri}}}{name}')
                            if el is None:
                                el = ET.SubElement(root, f'{{{uri}}}{name}')
                            el.text = 'de-DE' if name == 'language' else 'Klotzkette'
                        data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                    elif not has_core and entry.filename == '[Content_Types].xml':
                        root = ET.fromstring(data)
                        ET.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                                      PartName='/docProps/core.xml', ContentType='application/vnd.openxmlformats-package.core-properties+xml')
                        data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                    elif not has_core and entry.filename == '_rels/.rels':
                        root = ET.fromstring(data)
                        ET.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                                      Id='rIdKlotzketteCore', Type='http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties', Target='docProps/core.xml')
                        data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                    target.writestr(entry, data)
                if not has_core:
                    target.writestr('docProps/core.xml', '<?xml version="1.0" encoding="UTF-8"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:creator>Klotzkette</dc:creator><cp:lastModifiedBy>Klotzkette</cp:lastModifiedBy><dc:language>de-DE</dc:language></cp:coreProperties>')
            workbook.write_bytes(output.getvalue())


def load_script(filename, name):
    spec = importlib.util.spec_from_file_location(name, ROOT / 'scripts' / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def export():
    """Vorhandene Einzel- und ZIP-Builder nutzen; keine globalen Sammelarchive auslösen."""
    if not office_binary():
        raise RuntimeError('Native Office-Laufzeit fehlt. SOFFICE auf die konfigurierte Laufzeit setzen.')
    original = load_script('build-testakten-release-zips.py', 'bau_original')
    single = load_script('build-testakten-einzelpdf-zips.py', 'bau_single')
    for case in [WAR, BAD]:
        directory = ROOT / 'testakten' / case
        files = sorted(p for p in directory.iterdir() if p.is_file() and p.name[:2].isdigit())
        cache = single.render_office_batch([p for p in files if p.suffix in ('.docx','.xlsx')])
        writer = PdfWriter()
        rendered = {}
        for p in files:
            data = single.render_document_pdf(p, directory, cache)
            assert data
            data = metadata_pdf(data, p.stem)
            rendered[p] = data
            start = len(writer.pages)
            writer.append(PdfReader(io.BytesIO(data)))
            writer.add_outline_item(p.stem, start)
        writer.add_metadata({'/Author':'Klotzkette','/Title':case,'/Creator':'Klotzkette'})
        combined = directory / 'gesamt-pdf' / (case + '_gesamt.pdf')
        combined.parent.mkdir(exist_ok=True)
        writer.write(combined)
        shutil.copyfile(combined, ASSETS / combined.name)
        # Zentraler Builder erhält die native Office-Lesefassung aus dem verifizierten Cache.
        previous = single.render_document_pdf
        single.render_document_pdf = lambda p, d, office_cache=None: rendered[p]
        try:
            archive, count = single.build_single(directory, ASSETS)
        finally:
            single.render_document_pdf = previous
        original.build_single(directory, ASSETS)
        print(f'{case}: {len(files)} Originale, {len(writer.pages)} PDF-Seiten, {count} ZIP-Einträge')
        qa = QA / case / 'einzelpdf'
        qa.mkdir(parents=True, exist_ok=True)
        for p, data in rendered.items():
            (qa / (p.stem+'.pdf')).write_bytes(data)


def readmes():
    for case, title, desc in [(WAR, 'Werkhallenbau in Warendorf', 'Privater Werkhallenbau aus Sicht der Bauherrin und des Projektcontrollings. Kostenstand, Bestellungen, Restleistungen und Zahlungsplan treffen auf eine verspätete Trafostation und noch nicht freigegebene Zusatzangebote.'),
                              (BAD, 'Buchhaltungsprüfung in Bad Salzuflen', 'Begrenzter Rechnungs-, Leistungs- und Zahlungsabgleich eines regionalen Bauunternehmens. Der ausgewählte Stapel umfasst zwei Projekte, eine Lieferantenkorrektur, Skonto, einen Sicherheitseinbehalt und die im September gezahlten Augustlöhne.')]:
        directory = ROOT / 'testakten' / case
        files = sorted(p.name for p in directory.iterdir() if p.is_file() and p.name[:2].isdigit())
        text = f'<!-- decimal-headings -->\n\n# {title}\n\n## Vorgang\n\n{desc} Stand: 25. September 2026, 16:00 Uhr.\n\n<!-- BEGIN gesamt-pdf-section (autogen) -->\n## Downloads\n\nDie beiden Archive sind flach; das Originalarchiv enthält zusätzlich die Gesamtlesefassung. Der Formatmix bleibt im Originalarchiv erhalten. Die Einzel-PDF-Fassung enthält jede Unterlage als eigenes Dokument.\n\n'
        text += '> ' + NOTICE.replace('\n\n','\n>\n> ') + '\n\n'
        text += '| Fassung | Datei |\n| --- | --- |\n'
        text += f'| Gesamt-PDF | [Gesamtlesefassung](https://raw.githubusercontent.com/Klotzkette/claude-fuer-deutsches-recht/main/testakten/{case}/gesamt-pdf/{case}_gesamt.pdf) |\n'
        text += f'| Originale | [Akten-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{case}.zip) |\n'
        text += f'| Einzel-PDFs | [Einzel-PDF-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/latest/download/testakte-{case}-einzelpdfs.zip) |\n\n'
        text += '<!-- END gesamt-pdf-section (autogen) -->\n\n'
        text += f'## Bestand\n\n{len(files)} native Aktenstücke. Jede Datei bildet ein Dokument ab. Die beiden Arbeitsmappen sind bearbeitbar; E-Mail-Anhänge entsprechen den separat enthaltenen Quelldateien.\n\n| Nr. | Datei |\n| --- | --- |\n'
        text += '\n'.join(f'| {int(f[:2])} | `{f}` |' for f in files)
        text += '\n\n## Redaktion\n\nAutor: Klotzkette. Personen und Unternehmen des Sachverhalts sind erfunden. Die PNG-Dateien zeigen fachliche Bildschirmansichten des jeweiligen Falls. Die Prüfkriterien in `rubric.yaml` gehören nicht zum Export.\n'
        text += '\n## Quellenstand\n\nDie steuerlichen Fallannahmen wurden am 25. September 2026 mit den amtlichen Einzelnormen abgeglichen: '
        text += '[Paragraf 12 UStG](https://www.gesetze-im-internet.de/ustg_1980/__12.html) und [Paragraf 13b UStG](https://www.gesetze-im-internet.de/ustg_1980/__13b.html). '
        if case == BAD:
            text += 'Für den davon getrennten Bauabzug wurden [Paragraf 48 EStG](https://www.gesetze-im-internet.de/estg/__48.html) und [Paragraf 48b EStG](https://www.gesetze-im-internet.de/estg/__48b.html) geprüft. Die Bescheinigungsabschriften sind Fallunterlagen, keine tatsächlich erteilten Bescheinigungen. '
        else:
            text += 'Die Bauherrin fertigt Präzisionsteile und erbringt selbst keine Bauleistungen. Die Belege unterscheiden Nettoinvestitionen und Bruttozahlungen. '
        text += 'Rechtsprechung wird nicht verwendet.\n'
        (directory/'README.md').write_text(normalize_decimal_headings(text), encoding='utf-8')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--export-only', action='store_true')
    parser.add_argument('--check-runtime', action='store_true')
    args = parser.parse_args()
    QA.mkdir(parents=True, exist_ok=True)
    register_fonts()
    if args.check_runtime:
        assert office_binary(), 'SOFFICE fehlt'
        subprocess.run([node_binary(), str(ROOT/'scripts/build-bauwirtschaft-kaufmaennisch-workbooks.mjs'), '--check-runtime'], check=True)
        print('PDF-, Schrift-, Office- und Tabellenlaufzeiten verfügbar.')
        return
    if not args.export_only:
        warendorf()
        bad_salzuflen()
        payload = {'war':WAR, 'bad':BAD, 'root':str(ROOT), 'qa':str(QA), 'trades':TRADES,
                   'journal':W_JOURNAL, 'invoices':B_INVOICES, 'bank':B_BANK}
        data = QA / 'model.json'
        data.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
        subprocess.run([node_binary(), str(ROOT/'scripts/build-bauwirtschaft-kaufmaennisch-workbooks.mjs'), str(data)], check=True)
    office_metadata()
    readmes()
    export()


if __name__ == '__main__':
    main()
