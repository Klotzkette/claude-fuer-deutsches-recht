#!/usr/bin/env python3
"""Prüft die drei AML-Akten, Quellbelege, nativen Dateien und Releasepakete.

Ohne --dist werden die Quelldateien geprüft; nur die Paketprüfung entfällt.
Mit --dist VERZEICHNIS werden zusätzlich alle 19 Release-Assets geprüft.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
from functools import lru_cache
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import unittest
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader

from prompt_profiles import PROMPT_SUFFIXES, validate_files
from testakte_disclaimer import NOTICE_BYTES, pdf_content_errors
from testakte_einzelpdf_common import document_arcname_pairs
from testakte_zip_common import working_dump_archive_pairs

ROOT = Path(__file__).resolve().parents[1]
PLUGIN = ROOT / "geldwaeschebeauftragter"
DATA = json.loads((ROOT / "scripts/data/geldwaeschebeauftragter-akten.json").read_text(encoding="utf-8"))
VERSION = "445.34.0"
DIST = None
CATEGORIES = ("Vertragssoll", "Kassenbeleg", "Bankabgang", "Planwert", "Behauptung")
# Diese Erwartungen stehen bewusst unabhängig von der Klassifizierungsfunktion
# des Builders: Eine falsch zugeordnete Zahlung darf sich nicht selbst prüfen.
EXPECTED = {
    "aml-unternehmen-werkzeughandel-erfurt": {
        "transactions": [("U-T01", 39000, "Vertragssoll", "vereinbarter Gesamtpreis"),
                         ("U-T02", 6500, "Kassenbeleg", "belegt durch Kassenbeleg"),
                         ("U-T03", 6500, "Kassenbeleg", "belegt durch Kassenbeleg"),
                         ("U-T04", 26000, "Planwert", "angekündigt noch nicht gebucht")],
        "totals": (39000, 13000, "keine Angabe", 26000, "keine Angabe"),
        "counts": (1, 2, 0, 1, 0), "attachments": 6,
    },
    "aml-kanzlei-grundstueck-berlin": {
        "transactions": [("K-T01", 1280000, "Vertragssoll", "Kaufpreis im Entwurf"),
                         ("K-T02", 850000, "Planwert", "Finanzierungszusage mit Bedingungen"),
                         ("K-T03", 330000, "Planwert", "Eigenmittel angekündigt"),
                         ("K-T04", 100000, "Planwert", "Eigenmittel angekündigt")],
        "totals": (1280000, "keine Angabe", "keine Angabe", 1280000, "keine Angabe"),
        "counts": (1, 0, 0, 3, 0), "attachments": 8,
    },
    "aml-notariat-kaufpreis-wuerzburg": {
        "transactions": [("N-T01", 580000, "Vertragssoll", "beurkundeter Kaufpreis laut Fallunterlage"),
                         ("N-T02", 450000, "Bankabgang", "Bankabgang bestätigt Eingang noch nicht bestätigt"),
                         ("N-T03", 110000, "Bankabgang", "Bankabgang bestätigt Eingang noch nicht bestätigt"),
                         ("N-T04", 20000, "Behauptung", "Barzahlung nur widersprüchlich behauptet")],
        "totals": (580000, "keine Angabe", 560000, "keine Angabe", 20000),
        "counts": (1, 0, 2, 0, 1), "attachments": 8,
    },
}


def validator(filename):
    spec = importlib.util.spec_from_file_location(filename.replace("-", "_"), ROOT / "scripts" / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def case_directory(case):
    return ROOT / "testakten" / case["slug"]


def pdf_text(data):
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)


def normalized(text):
    return re.sub(r"\s+", " ", text).strip()


def normalized_pdf_passage(text):
    # LibreOffice darf ein Aktenzeichen nach einem Bindestrich umbrechen.
    # Nur den dabei extrahierten Leerraum glätten; alle Textzeichen behalten.
    return normalized(re.sub(r"(?<=\w)-\s+(?=\w)", "-", text))


@lru_cache(maxsize=None)
def source_text(path):
    if path.suffix == ".pdf":
        return pdf_text(path.read_bytes())
    if path.suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            tree = ElementTree.fromstring(archive.read("word/document.xml"))
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        return "\n".join("".join(p.itertext()) for p in tree.iter(ns + "p"))
    if path.suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        return str(message["Message-ID"]) + "\n" + message.get_body(preferencelist=("plain",)).get_content()
    raise ValueError(f"Nicht unterstützter Quellenpfad: {path}")


def release_copies():
    paths = [PLUGIN / f"geldwaeschebeauftragter-{kind}.{ext}"
             for kind in ("werkstatt", "schnellstart", "hauptproblem") for ext in ("md", "txt")]
    paths += [case_directory(case) / "gesamt-pdf" / f"{case['slug']}_gesamt.pdf" for case in DATA["cases"]]
    paths += [ROOT / "docs/handbuecher" / name for name in
              ("geldwaeschebeauftragter-skills-handbuch.pdf", "geldwaeschebeauftragter-werkstatt-lesefassung.pdf")]
    return paths


class GeldwaeschebeauftragterTest(unittest.TestCase):
    def test_pdf_passage_normalization_preserves_identifiers(self):
        expected = "Vorgang NF-2026-0915: Kaufpreis 580.000,00 EUR."
        for wrapped in ("NF-\n2026-0915", "NF- \n2026-0915", "NF- 2026-0915"):
            self.assertEqual(normalized_pdf_passage(expected.replace("NF-2026-0915", wrapped)), expected)
        for changed in ("NF-2026-0916", "NF2026-0915", "NF-2026-095", "NF-2026--0915"):
            self.assertNotEqual(normalized_pdf_passage(expected.replace("NF-2026-0915", changed)), expected)
        self.assertEqual(normalized_pdf_passage("Kaufpreis - 580.000,00 EUR"), "Kaufpreis - 580.000,00 EUR")
        self.assertEqual(normalized_pdf_passage("580.000,00 EUR"), "580.000,00 EUR")

    def test_canonical_inventory_and_original_formats(self):
        self.assertEqual(DATA["schema_version"], 1)
        self.assertTrue(DATA["fictional"])
        self.assertEqual(DATA["reference_date"], "2026-10-08")
        self.assertEqual(DATA["currency"], "EUR")
        self.assertEqual(Counter(c["slug"] for c in DATA["cases"]), Counter(EXPECTED.keys()))
        identifiers = []
        for case in DATA["cases"]:
            with self.subTest(case=case["slug"]):
                directory = case_directory(case)
                declared = [d["path"] for d in case["documents"] + case["emails"]]
                declared.append("04_Tabellen/Fallregister.xlsx")
                self.assertEqual(len(declared), 25)
                self.assertEqual(len(set(declared)), 25)
                self.assertEqual(Counter(Path(p).suffix for p in declared),
                                 {".pdf": 8, ".docx": 4, ".eml": 12, ".xlsx": 1})
                actual = [p.relative_to(directory).as_posix() for p in directory.rglob("*")
                          if p.is_file() and p.suffix.lower() in {".pdf", ".docx", ".eml", ".xlsx"}
                          and "gesamt-pdf" not in p.relative_to(directory).parts]
                self.assertCountEqual(actual, declared)
                # Die zentrale Exportauswahl darf keinen als Rohunterlage
                # bestellten Bestandteil als redaktionelle Metadaten verlieren.
                selected = [p.relative_to(directory).as_posix() for p, _ in
                            working_dump_archive_pairs(directory, include_gesamt_pdf=False)]
                self.assertCountEqual(selected, declared)
                self.assertEqual(len(document_arcname_pairs(directory)), 25)
                for relative in declared:
                    self.assertGreater((directory / relative).stat().st_size, 0, relative)
                identifiers.extend(d["id"] for d in case["documents"] + case["emails"] + case["transactions"])
        self.assertEqual(len(identifiers), len(set(identifiers)))

    def test_source_documents_reproduce_canonical_content(self):
        for case in DATA["cases"]:
            for document in case["documents"]:
                path = case_directory(case) / document["path"]
                with self.subTest(source=path.name):
                    text = normalized(source_text(path))
                    self.assertTrue(normalized(document["title"]) in text, f"Titel fehlt in {path.name}")
                    for paragraph in document["paragraphs"]:
                        self.assertTrue(normalized(paragraph) in text, f"Quellabsatz fehlt in {path.name}: {paragraph[:100]}")
                    for row in document.get("table", []):
                        for value in row:
                            self.assertIn(normalized(str(value)), text)
                    if path.suffix == ".pdf":
                        self.assertIn(document["id"], text)
                        self.assertEqual(pdf_content_errors(path.read_bytes()), [])

    def test_evidence_ids_resolve_to_actual_source_documents(self):
        for case in DATA["cases"]:
            sources = {d["id"]: d for d in case["documents"] + case["emails"]}
            for entry in case["transactions"] + case["ownership"] + case["chronology"]:
                identifier = entry["evidence_id"]
                with self.subTest(case=case["slug"], evidence=identifier):
                    self.assertIn(identifier, sources)
                    path = case_directory(case) / sources[identifier]["path"]
                    self.assertTrue(path.is_file())
                    self.assertIn(identifier.lower(), source_text(path).lower())
            for transaction in case["transactions"]:
                self.assertIs(type(transaction["amount_eur"]), int)
                self.assertGreater(transaction["amount_eur"], 0)
                date.fromisoformat(transaction["date"])

    def test_eml_headers_threads_reserved_contacts_and_binary_attachments(self):
        message_ids, total_attachments = set(), 0
        for case in DATA["cases"]:
            contacts = {c["id"]: c for c in case["contacts"]}
            documents = {d["id"]: d for d in case["documents"]}
            emails = {e["id"]: e for e in case["emails"]}
            attached = 0
            for contact in contacts.values():
                self.assertTrue(contact["email"].rsplit("@", 1)[-1].endswith(".example"))
            for entry in case["emails"]:
                path = case_directory(case) / entry["path"]
                with self.subTest(email=path.name):
                    msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                    for header in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version"):
                        self.assertTrue(msg[header], f"{path}: {header}")
                    self.assertTrue(all(not part.defects for part in msg.walk()))
                    identifier = str(msg["Message-ID"])
                    self.assertNotIn(identifier, message_ids)
                    message_ids.add(identifier)
                    self.assertEqual(identifier, f"<{entry['id'].lower()}-{entry['date']}@aktenpost.example>")
                    self.assertEqual(str(msg["Subject"]), entry["subject"])
                    stamp = parsedate_to_datetime(str(msg["Date"]))
                    self.assertEqual(stamp.date().isoformat(), entry["date"])
                    self.assertIsNotNone(stamp.utcoffset())
                    self.assertEqual([a for _, a in getaddresses(msg.get_all("From", []))],
                                     [contacts[entry["sender"]]["email"]])
                    self.assertEqual([a for _, a in getaddresses(msg.get_all("To", []))],
                                     [contacts[c]["email"] for c in entry["recipients"]])
                    for _, address in getaddresses(msg.get_all("From", []) + msg.get_all("To", [])):
                        self.assertTrue(address.rsplit("@", 1)[-1].endswith(".example"))
                    self.assertTrue(msg.get_body(preferencelist=("plain",)).get_content().startswith(entry["body"]))
                    if entry.get("reply_to"):
                        previous = emails[entry["reply_to"]]
                        reference = f"<{previous['id'].lower()}-{previous['date']}@aktenpost.example>"
                        self.assertEqual(str(msg["In-Reply-To"]), reference)
                        self.assertEqual(str(msg["References"]), reference)
                    else:
                        self.assertIsNone(msg["In-Reply-To"])
                    parts = list(msg.iter_attachments())
                    sources = [case_directory(case) / documents[key]["path"] for key in entry["attachments"]]
                    self.assertEqual([part.get_filename() for part in parts], [p.name for p in sources])
                    for part, source in zip(parts, sources):
                        self.assertEqual(part.get_payload(decode=True), source.read_bytes(), source.name)
                    attached += len(parts)
            self.assertEqual(attached, EXPECTED[case["slug"]]["attachments"])
            total_attachments += attached
        self.assertEqual(len(message_ids), 36)
        self.assertEqual(total_attachments, 22)

    def test_native_xlsx_rows_formula_caches_and_evidence_status(self):
        for case in DATA["cases"]:
            path = case_directory(case) / "04_Tabellen/Fallregister.xlsx"
            expected = EXPECTED[case["slug"]]
            with self.subTest(case=case["slug"]):
                formulas = load_workbook(path, data_only=False)
                cached = load_workbook(path, data_only=True)
                self.assertEqual(formulas.sheetnames, ["Übersicht", "Zahlungen", "Beteiligte", "Verlauf"])
                self.assertEqual(len(case["transactions"]), len(expected["transactions"]))
                for row_number, (source, wanted) in enumerate(zip(case["transactions"], expected["transactions"]), 7):
                    identifier, amount, category, status = wanted
                    self.assertEqual((source["id"], source["amount_eur"], source["status"]), (identifier, amount, status))
                    row = [cached["Zahlungen"].cell(row_number, col).value for col in range(1, 10)]
                    self.assertEqual(row[0], identifier)
                    self.assertEqual(row[1].date().isoformat(), source["date"])
                    self.assertEqual(row[2:], [amount, category, status, source["payer"], source["payee"],
                                              source["evidence_id"], source["allocation"]])
                summary = cached["Übersicht"]
                self.assertEqual(tuple(summary[f"A{r}"].value for r in range(8, 13)), CATEGORIES)
                self.assertEqual(tuple(summary[f"B{r}"].value for r in range(8, 13)), expected["totals"])
                self.assertEqual(tuple(summary[f"C{r}"].value for r in range(8, 13)), expected["counts"])
                self.assertIn("Eingang beim Empfänger ist damit noch nicht bestätigt", summary["D10"].value)
                self.assertIn("nicht als erfolgte Zahlung", summary["D11"].value)
                self.assertIn("Keine bestätigte Kaufpreisleistung", summary["D12"].value)
                self.assertIn("Betragsarten nicht addieren", summary["A34"].value)
                formula_cells = []
                for sheet in formulas:
                    for row in sheet:
                        for cell in row:
                            value = cached[sheet.title][cell.coordinate]
                            self.assertNotEqual(value.data_type, "e", f"{path}/{sheet.title}/{cell.coordinate}")
                            if cell.data_type == "f":
                                formula_cells.append((sheet.title, cell.coordinate))
                                self.assertIsNotNone(value.value, f"Formelcache fehlt: {path}/{sheet.title}/{cell.coordinate}")
                self.assertCountEqual(formula_cells, [("Übersicht", f"{col}{row}") for col in "BC" for row in range(8, 13)])
                end = 6 + len(case["transactions"])
                for row in range(8, 13):
                    self.assertEqual(formulas["Übersicht"][f"C{row}"].value,
                                     f"=COUNTIFS('Zahlungen'!$D$7:$D${end},A{row})")
                    self.assertEqual(formulas["Übersicht"][f"B{row}"].value,
                                     f'=IF(C{row}=0,"keine Angabe",SUMIFS(\'Zahlungen\'!$C$7:$C${end},\'Zahlungen\'!$D$7:$D${end},A{row}))')
                formulas.close()
                cached.close()

    def test_skills_prompts_and_versions(self):
        skills = list((PLUGIN / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 11)
        for skill in skills:
            with self.subTest(skill=skill.parent.name):
                body = skill.read_text(encoding="utf-8")
                self.assertEqual(re.findall(r"^## ([1-6])\. ", body, re.M), list("123456"))
                frontmatter = body.split("---", 2)[1]
                self.assertEqual(re.findall(r"^([a-z_]+):", frontmatter, re.M), ["name", "description"])
                self.assertIn("name: " + skill.parent.name, frontmatter)
                self.assertLessEqual(len(skill.parent.name), 64)
                self.assertTrue(skill.parent.name.isascii())
                description = re.search(r"^description: (.*)$", frontmatter, re.M).group(1)
                self.assertLessEqual(len(description), 360)
                output = body.split("## 5.", 1)[1].split("## 6.", 1)[0]
                for required in ("Times New Roman", "11 pt", "ausformuliert"):
                    self.assertIn(required, output)
                self.assertTrue("references/zitierweise.md" in body, f"Zitierweise-Verweis fehlt: {skill.parent.name}")
        self.assertEqual(validate_files(PLUGIN, PLUGIN.name, ROOT), [])
        for kind in ("werkstatt", "schnellstart", "hauptproblem"):
            md = PLUGIN / f"geldwaeschebeauftragter-{kind}.md"
            self.assertEqual(md.read_bytes(), md.with_suffix(".txt").read_bytes())
            self.assertGreater(len(md.read_text(encoding="utf-8")), 1000)
            if kind != "werkstatt":
                self.assertLessEqual(len(md.read_bytes()), 7500)
        for directory in (".claude-plugin", ".codex-plugin"):
            manifest = json.loads((PLUGIN / directory / "plugin.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["name"], PLUGIN.name)
            self.assertEqual(manifest["version"], VERSION)
        market = json.loads((ROOT / ".claude-plugin/marketplace.json").read_text(encoding="utf-8"))
        entry = [p for p in market["plugins"] if p["name"] == PLUGIN.name]
        self.assertEqual(len(entry), 1)
        self.assertEqual(entry[0]["version"], VERSION)
        routes = json.loads((ROOT / "scripts/scoped-release-assets.json").read_text(encoding="utf-8"))["assets"]
        self.assertEqual(routes["geldwaeschebeauftragter.zip"], f"geldwaeschebeauftragter-v{VERSION}")

    def test_relative_markdown_links_resolve(self):
        files = list(PLUGIN.rglob("*.md")) + [case_directory(case) / "README.md" for case in DATA["cases"]]
        for path in files:
            for target in re.findall(r"\[[^\]]*\]\(([^)]+)\)", path.read_text(encoding="utf-8")):
                parsed = urlsplit(target.strip("<>"))
                if parsed.scheme or parsed.netloc or not parsed.path:
                    continue
                with self.subTest(file=path.relative_to(ROOT), target=target):
                    self.assertTrue((path.parent / unquote(parsed.path)).resolve().exists())

    def test_combined_pdfs_keep_source_ids_without_notice_pages(self):
        for case in DATA["cases"]:
            path = case_directory(case) / "gesamt-pdf" / f"{case['slug']}_gesamt.pdf"
            with self.subTest(case=case["slug"]):
                self.assertTrue(path.is_file(), f"Gesamt-PDF fehlt: {path}")
                data = path.read_bytes()
                self.assertEqual(pdf_content_errors(data), [])
                text = pdf_text(data).lower()
                for source in case["documents"] + case["emails"]:
                    self.assertTrue(source["id"].lower() in text, f"{path.name}: Quellen-ID fehlt: {source['id']}")
                for transaction in case["transactions"]:
                    self.assertTrue(transaction["id"].lower() in text, f"{path.name}: Vorgang-ID fehlt: {transaction['id']}")

    def test_release_packages_and_checksums(self):
        if DIST is None:
            self.skipTest("Paketprüfung nur mit --dist VERZEICHNIS; Quelldateien werden separat geprüft")
        originals = validator("validate-testakten-release-zips.py")
        singles = validator("validate-testakten-einzelpdf-zips.py")
        expected = {"geldwaeschebeauftragter.zip", "checksums-sha256.txt", *(p.name for p in release_copies())}
        expected.update(f"testakte-{slug}{suffix}.zip" for slug in EXPECTED for suffix in ("", "-einzelpdfs"))
        self.assertEqual(len(expected), 19)
        self.assertEqual({p.name for p in DIST.iterdir()}, expected)
        for source in release_copies():
            self.assertEqual((DIST / source.name).read_bytes(), source.read_bytes(), source.name)
        for case in DATA["cases"]:
            directory = case_directory(case)
            original_zip = DIST / f"testakte-{case['slug']}.zip"
            single_zip = DIST / f"testakte-{case['slug']}-einzelpdfs.zip"
            # Unveränderte zentrale Kernvalidatoren: keine Sonderbefreiung für
            # Warntext, A4-Format, Namen, doppelte Einträge oder Archivstruktur.
            originals.assert_same(original_zip.name, originals.expected_entries(directory),
                                  originals.zip_entries(original_zip, require_notice=True))
            singles.assert_same(single_zip.name, ["README.txt", *singles.expected_arcnames(directory)],
                                singles.zip_entries(single_zip, expected_suffix=".pdf"))
            for archive_path in (original_zip, single_zip):
                with zipfile.ZipFile(archive_path) as archive:
                    self.assertEqual(archive.namelist()[0], "README.txt")
                    self.assertEqual(archive.read("README.txt"), NOTICE_BYTES)
                    self.assertFalse(any("/" in n or n.endswith(".md") for n in archive.namelist()))
            with zipfile.ZipFile(original_zip) as archive:
                pairs = working_dump_archive_pairs(directory, include_gesamt_pdf=True)
                self.assertEqual(len(pairs), 26)
                for source, name in pairs:
                    self.assertEqual(archive.read(name), source.read_bytes(), name)
            with zipfile.ZipFile(single_zip) as archive:
                pairs = document_arcname_pairs(directory)
                self.assertEqual(len(pairs), 25)
                for source, name in pairs:
                    data = archive.read(name)
                    if source.suffix == ".pdf":
                        self.assertEqual(data, source.read_bytes(), name)
                    text = pdf_text(data).lower()
                    if source.suffix == ".xlsx":
                        for transaction in case["transactions"]:
                            self.assertTrue(transaction["id"].lower() in text, f"{name}: Vorgang-ID fehlt: {transaction['id']}")
                    elif source.suffix == ".docx":
                        original = next(d for d in case["documents"] if d["path"] == source.relative_to(directory).as_posix())
                        for passage in [original["title"], *original["paragraphs"]]:
                            self.assertIn(normalized_pdf_passage(passage.lower()), normalized_pdf_passage(text), f"{name}: Word-Inhalt fehlt")
                    else:
                        self.assertTrue(source.name.split("_", 1)[0].lower() in text, f"{name}: Quellen-ID fehlt")
        with zipfile.ZipFile(DIST / "geldwaeschebeauftragter.zip") as archive:
            expected_plugin = {p.relative_to(PLUGIN).as_posix() for p in PLUGIN.rglob("*")
                               if p.is_file() and not p.is_symlink() and "__pycache__" not in p.parts
                               and p.suffix != ".pyc" and p.name not in {".DS_Store", "CLAUDE.md"}
                               and not p.name.endswith(PROMPT_SUFFIXES)}
            self.assertCountEqual(archive.namelist(), expected_plugin)
            self.assertIsNone(archive.testzip())
            self.assertEqual(len([n for n in archive.namelist() if n.endswith("/SKILL.md")]), 11)
            for name in archive.namelist():
                self.assertFalse(name.startswith("testakten/") or "__pycache__" in name or name.endswith(PROMPT_SUFFIXES))
                self.assertEqual(archive.read(name), (PLUGIN / name).read_bytes(), name)
        checksum_names = []
        for line in (DIST / "checksums-sha256.txt").read_text(encoding="utf-8").splitlines():
            self.assertRegex(line, r"^[0-9a-f]{64}  [^/\\]+$")
            digest, name = line.split("  ", 1)
            checksum_names.append(name)
            self.assertEqual(hashlib.sha256((DIST / name).read_bytes()).hexdigest(), digest, name)
        self.assertCountEqual(checksum_names, expected - {"checksums-sha256.txt"})


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dist", type=Path, help="Optionales Verzeichnis mit den 19 Release-Assets")
    args, remaining = parser.parse_known_args()
    DIST = args.dist.resolve() if args.dist else None
    unittest.main(argv=[__file__, *remaining])
