#!/usr/bin/env python3
"""Prüft Quellenpaket, eigenständige Prompts und bytegenaue Release-Zuordnung."""
from __future__ import annotations
import argparse
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import unittest
from urllib.parse import unquote
import zipfile
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import yaml
from quality_lab import validate_profile

ROOT = Path(__file__).resolve().parents[1]
NAME = 'insolvenzanfechtung'
PARENT = 'fachanwalt-insolvenz-sanierungsrecht'
DIST = None
VEEDELWENDE_CASES = (
    'insolvenzanfechtung-veedelwende-direktrueckzahlung-koeln',
    'insolvenzanfechtung-veedelwende-zession-koeln',
)


def module(filename):
    spec = importlib.util.spec_from_file_location(filename.replace('-', '_'), ROOT/'scripts'/filename)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def compact(text):
    return re.sub(r'\s+', ' ', text).strip()


def document_text(path):
    if path.suffix == '.docx':
        return '\n'.join(paragraph.text for paragraph in Document(path).paragraphs)
    if path.suffix == '.pdf':
        return '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages)
    if path.suffix == '.eml':
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        body = message.get_body(preferencelist=('plain',))
        if body is None:
            raise AssertionError(f'Keine Textfassung in {path}')
        return body.get_content()
    if path.suffix == '.txt':
        return path.read_text(encoding='utf-8')
    if path.suffix == '.xlsx':
        workbook = load_workbook(path, data_only=False, read_only=True)
        try:
            return '\n'.join(
                str(cell.value)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
        finally:
            workbook.close()
    raise ValueError(f'Kein Textleser für {path}')


def veedelwende_cases():
    cases = json.loads((ROOT/'scripts/data/insolvenzanfechtung-akten.json').read_text(encoding='utf-8'))
    return {case['slug']: case for case in cases if case['slug'] in VEEDELWENDE_CASES}


def case_items(case):
    return {item['name']: item for item in case['files']}


class AnfechtungTest(unittest.TestCase):
    def test_eleven_skills_are_activatable_and_references_resolve(self):
        skills = sorted((ROOT/NAME/'skills').glob('*/SKILL.md'))
        self.assertEqual(len(skills), 11)
        for path in skills:
            text = path.read_text(encoding='utf-8')
            front = yaml.safe_load(text.split('---', 2)[1])
            self.assertEqual(set(front), {'name', 'description'})
            self.assertEqual(front['name'], path.parent.name)
            self.assertLessEqual(len(front['description']), 360)
            self.assertNotIn('§', front['description'])
            self.assertNotIn('im Plugin', front['description'])
            self.assertEqual(re.findall(r'^## ([1-6])\.', text, re.M), list('123456'))
            for word in ('Times New Roman', '11 pt', 'ausformuliert'):
                self.assertIn(word, text)
            for target in re.findall(r'\]\(([^)]+)\)', text):
                if target.startswith(('https:', 'http:', '#', 'mailto:')):
                    continue
                target = unquote(target.split('#')[0])
                self.assertTrue((path.parent/target).resolve().exists(), f'{path}: {target}')

    def test_prompts_are_separate_complete_and_bounded(self):
        builder = module('build-insolvenzanfechtung-release.py')
        for source in builder.prompt_sources():
            self.assertEqual(source.read_bytes(), source.with_suffix('.txt').read_bytes())
            data = source.read_bytes().replace(b'\r\n', b'\n').replace(b'\n', b'\r\n')
            if 'werkstatt' not in source.name:
                self.assertLessEqual(len(data), 7500, source)
            else:
                self.assertGreater(len(source.read_text().split()), 6000)
            self.assertIn('Trägt nicht', source.read_text())
        profile = json.loads((ROOT/'quality/evals'/f'{NAME}.json').read_text())
        validate_profile(profile, NAME, ROOT/NAME)
        self.assertEqual(len(profile['sources']), 14)
        self.assertEqual(len(profile['cases']), 5)

    def test_section_135_and_assignment_anchors_are_present(self):
        skill = (ROOT/NAME/'skills/gesellschafterfinanzierung/SKILL.md').read_text()
        workshop = (ROOT/NAME/f'{NAME}-werkstatt.md').read_text()
        sources = (ROOT/NAME/'references/rechtsquellen.md').read_text()
        for anchor in ('IX ZR 196/13', 'IX ZR 32/12', 'IX ZR 173/23'):
            self.assertIn(anchor, skill)
            self.assertIn(anchor, workshop)
            self.assertIn(anchor, sources)
        self.assertIn('objektive Gläubigerbenachteiligung', skill)
        self.assertIn('keine zeitlich unbegrenzte', sources)
        self.assertIn('Treuhand', workshop)

    def test_veedelwende_case_files_match_the_scoped_data(self):
        cases = veedelwende_cases()
        self.assertEqual(set(cases), set(VEEDELWENDE_CASES))
        original_suffixes = {'.docx', '.pdf', '.xlsx', '.eml', '.txt'}
        for slug, case in cases.items():
            items = case['files']
            self.assertEqual(len(items), 24, slug)
            expected = {item['name'] for item in items}
            self.assertEqual(len(expected), 24, slug)
            directory = ROOT/'testakten'/slug
            actual = {path.name for path in directory.iterdir()
                      if path.is_file() and path.suffix in original_suffixes}
            self.assertEqual(actual, expected, slug)
            for item in items:
                path = directory/item['name']
                self.assertGreater(path.stat().st_size, 0, path)
                if item.get('body'):
                    self.assertIn(compact(item['body']), compact(document_text(path)), path)
                elif path.suffix == '.xlsx':
                    text = document_text(path)
                    self.assertIn(item['title'], text, path)
                    for value in (value for row in item['rows'] for value in row):
                        self.assertIn(str(value), text, path)

    def test_veedelwende_direct_case_keeps_exact_facts_and_clean_complaint(self):
        case = veedelwende_cases()[VEEDELWENDE_CASES[0]]
        items = case_items(case)
        self.assertEqual(
            {name: items[name]['date'] for name in (
                '02_Insolvenzantrag_Eingang.pdf',
                '05_Darlehensvertrag_Haderlein.docx',
                '13_Kontoauszug_Rueckzahlung.pdf',
                '21_Klageschrift.docx',
            )},
            {
                '02_Insolvenzantrag_Eingang.pdf': '03.06.2026',
                '05_Darlehensvertrag_Haderlein.docx': '12.06.2024',
                '13_Kontoauszug_Rueckzahlung.pdf': '04.02.2026',
                '21_Klageschrift.docx': '15.09.2026',
            },
        )
        directory = ROOT/'testakten'/case['slug']
        for filename, facts in {
            '02_Insolvenzantrag_Eingang.pdf': ('3. Juni 2026',),
            '05_Darlehensvertrag_Haderlein.docx': ('150.000,00 EUR',),
            '13_Kontoauszug_Rueckzahlung.pdf': ('150.000,00 EUR', '3. Februar 2026'),
        }.items():
            text = document_text(directory/filename)
            for fact in facts:
                self.assertIn(fact, text, filename)

        complaint = document_text(directory/'21_Klageschrift.docx')
        for fact in ('150.000,00 EUR', '3. Februar 2026', '3. Juni 2026',
                     'IX ZR 196/13', 'IX ZR 44/22', 'Nala Seifert'):
            self.assertIn(fact, complaint)
        self.assertNotRegex(complaint, r'(?i)\bTilda\b|\b40\s*k\b|\b40[.\s]?000(?:,00)?\b')

    def test_veedelwende_assignment_case_exposes_timing_and_litigation_risk(self):
        case = veedelwende_cases()[VEEDELWENDE_CASES[1]]
        items = case_items(case)
        self.assertEqual(items['07_Kauf_Abtretung.docx']['date'], '03.12.2024')
        self.assertEqual(items['08_Kaufpreiszahlung_Quast.pdf']['date'], '06.12.2024')
        self.assertEqual(items['09_Abtretungsanzeige.eml']['date'], '2024-12-07T10:17:00+01:00')
        purchase_date = datetime.strptime(items['08_Kaufpreiszahlung_Quast.pdf']['date'], '%d.%m.%Y').date()
        notice_date = datetime.fromisoformat(items['09_Abtretungsanzeige.eml']['date']).date()
        self.assertEqual((purchase_date, notice_date), (date(2024, 12, 6), date(2024, 12, 7)))
        self.assertLess(purchase_date, notice_date)

        directory = ROOT/'testakten'/case['slug']
        purchase = document_text(directory/'08_Kaufpreiszahlung_Quast.pdf')
        self.assertIn('48.000,00 EUR', purchase)
        self.assertIn('6. Dezember 2024', purchase)
        notice = BytesParser(policy=policy.default).parsebytes(
            (directory/'09_Abtretungsanzeige.eml').read_bytes()
        )
        self.assertEqual(parsedate_to_datetime(str(notice['Date'])).date(), date(2024, 12, 7))
        repayment = document_text(directory/'12_Kontoauszug_Rueckzahlung.pdf')
        for fact in ('96.000,00 EUR', '3. April 2026', '3. Juni 2026'):
            self.assertIn(fact, repayment)

        texts = {
            'Risikomemo': document_text(directory/'21_Risikomemo_BGH.docx'),
            'Klage': document_text(directory/'22_Klageschrift.docx'),
            'Klageerwiderung': document_text(directory/'24_Klageerwiderung.docx'),
        }
        for label, text in texts.items():
            for fact in ('IX ZR 32/12', 'IX ZR 173/23', '18 Monate', '17. April 2026'):
                self.assertIn(fact, text, label)
            self.assertRegex(text, r'(?i)(vier Wochen|vierwöchig|Vierwochenfrist)', label)
        self.assertIn('offen risikobehafteter Versuch', texts['Risikomemo'])
        self.assertIn('Ohne neue Tatsachen', texts['Risikomemo'])
        self.assertIn('erhebliches Klagerisiko', texts['Klage'])
        self.assertIn('mit einer Klageabweisung zu rechnen', texts['Klage'])
        self.assertIn('Die Klage ist abzuweisen', texts['Klageerwiderung'])

    def test_veedelwende_email_attachments_are_byte_exact(self):
        for slug, case in veedelwende_cases().items():
            directory = ROOT/'testakten'/slug
            for item in case['files']:
                if not item.get('attachments'):
                    continue
                message = BytesParser(policy=policy.default).parsebytes(
                    (directory/item['name']).read_bytes()
                )
                attachments = [
                    (part.get_filename(), part.get_payload(decode=True))
                    for part in message.iter_attachments()
                ]
                self.assertEqual([name for name, _ in attachments], item['attachments'], item['name'])
                for name, payload in attachments:
                    self.assertEqual(payload, (directory/name).read_bytes(), f'{item["name"]}: {name}')

    def test_manifests_and_marketplace_agree(self):
        market = json.loads((ROOT/'.claude-plugin/marketplace.json').read_text())
        for name, version in ((NAME, '1.1.0'), (PARENT, '445.35.4')):
            entries = [p for p in market['plugins'] if p['name'] == name]
            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0]['version'], version)
            for relative in (('plugin.json', '.claude-plugin/plugin.json', '.codex-plugin/plugin.json') if name == NAME else ('.claude-plugin/plugin.json',)):
                manifest = json.loads((ROOT/name/relative).read_text())
                self.assertEqual((manifest['name'], manifest['version']), (name, version))
                self.assertEqual(manifest['description'], entries[0]['description'])

    def test_release_assets_and_packages_are_byte_exact(self):
        if DIST is None:
            self.skipTest('--dist erforderlich')
        builder = module('build-insolvenzanfechtung-release.py')
        self.assertEqual({p.name for p in DIST.iterdir()}, builder.asset_names())
        lines = (DIST/'checksums-sha256.txt').read_text().splitlines()
        self.assertEqual(len(lines), len(builder.asset_names())-1)
        for line in lines:
            digest, name = line.split('  ', 1)
            self.assertEqual(sha(DIST/name), digest)
        for name in (NAME, PARENT):
            expected = builder.plugin_sources(name)
            for suffix, prefix in (('', ''), ('-portable', name+'/')):
                with zipfile.ZipFile(DIST/(name+suffix+'.zip')) as archive:
                    self.assertIsNone(archive.testzip())
                    self.assertEqual(set(archive.namelist()), {prefix+p for p in expected})
                    for relative, path in expected.items():
                        self.assertEqual(archive.read(prefix+relative), path.read_bytes())
                    for path in archive.namelist():
                        self.assertNotIn('..', Path(path).parts)
                        self.assertFalse(set(Path(path).parts) & {'testakte', 'testakten', 'prompts'})
        record = json.loads((DIST/'quellenabgleich.json').read_text())
        for source in builder.prompt_sources():
            asset = builder.prompt_asset(source)
            self.assertEqual((DIST/asset).read_bytes(), source.read_bytes())
            target = DIST/str(Path(asset).with_suffix('.pdf'))
            result = builder.inspect_pdf(source, target)
            for key, value in result.items():
                self.assertEqual(record['pdfs'][target.name][key], value)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dist', type=Path)
    args, remaining = parser.parse_known_args()
    DIST = args.dist.resolve() if args.dist else None
    unittest.main(argv=[__file__, *remaining])
