#!/usr/bin/env python3
"""Erzeugt nur den getrennten Nachtrag der drei Gesellschafterstreit-Akten.

Kein Release, keine Gesamt-PDF, keine Änderung der Kernunterlagen. Die drei
Fall-READMEs erhalten einen idempotenten Nachtragsabschnitt. Tabellen werden
mit dem gleichnamigen MJS-Builder erzeugt. Gebündelte Laufzeit erforderlich.
"""
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.headerregistry import Address
from email.utils import getaddresses
import argparse
import hashlib
import importlib.util
import io
import json
import os
import re
import subprocess
import sys
import zipfile
from concurrent.futures import ThreadPoolExecutor
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject, ArrayObject, ByteStringObject
from gesellschafterstreit_nachtrag_daten import CASES

ROOT = Path(__file__).resolve().parents[1]
SUBDIR = 'Nachtrag_2026-10-08'
RUNTIME = Path('/Users/klotzkette/.cache/codex-runtimes/codex-primary-runtime')
RENDER = Path('/Users/klotzkette/.codex/plugins/cache/openai-primary-runtime/documents/26.1007.11041/skills/documents/render_docx.py')
WARNING = '> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n>\n> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'

def word(item, out):
    d = Document(); sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.25)
    sec.top_margin = sec.bottom_margin = Cm(1.9)
    sec.header_distance = sec.footer_distance = Cm(1)
    for style in d.styles:
        if style.type == 1:
            style.font.name = 'Times New Roman'
            style.font.size = Pt(11)
            style.font.color.rgb = RGBColor(0,0,0)
    norm = d.styles['Normal'].paragraph_format
    norm.line_spacing = 1.08; norm.space_after = Pt(7); norm.widow_control = True
    for name, size in [('Title',15),('Heading 1',11)]:
        s = d.styles[name]; s.font.size = Pt(size); s.font.bold = True
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(7)
        s.paragraph_format.space_after = Pt(7)
    for node in list(d.styles.element.iter()):
        if node.tag == qn('w:pBdr'): node.getparent().remove(node)
        if node.tag == qn('w:rFonts'):
            for key in list(node.attrib):
                if 'theme' in key.lower(): del node.attrib[key]
            node.set(qn('w:ascii'),'Times New Roman'); node.set(qn('w:hAnsi'),'Times New Roman')
    d.add_paragraph(item['sender'])
    d.add_paragraph('An: '+item['recipient'])
    d.add_paragraph(item['date'])
    d.add_paragraph(item['title'],'Title')
    for block in re.split(r'\n\s*\n',item['body']):
        if block.startswith('## '): d.add_paragraph(block[3:].strip(),'Heading 1')
        else: d.add_paragraph(block)
    f = sec.footer.paragraphs[0]; f.alignment = 2
    f.add_run('Seite ').font.size=Pt(11)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
    d.core_properties.title=item['title']; d.core_properties.author=item['sender'].splitlines()[0]
    d.core_properties.created=d.core_properties.modified=datetime(2026,10,8,12)
    buffer=io.BytesIO(); d.save(buffer)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as z, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as target:
        for name in sorted(z.namelist()):
            info=zipfile.ZipInfo(name,(2026,10,8,12,0,0)); info.compress_type=zipfile.ZIP_DEFLATED
            target.writestr(info,z.read(name))

def message(item,directory):
    m=EmailMessage(policy=SMTP)
    for header,key in [('From','sender'),('To','recipient')]:
        m[header]=tuple(Address(display_name=name,addr_spec=address) for name,address in getaddresses([item[key]]))
    m['Subject']=item['title']
    m['Date']=datetime.fromisoformat(item['date'])
    ident=hashlib.sha256((directory.parent.name+item['file']).encode()).hexdigest()[:24]
    m['Message-ID']=f'<{ident}@akten-nachtrag.example>'
    m.set_content(item['body']+'\n',charset='utf-8')
    types={'.pdf':('application','pdf'),'.docx':('application','vnd.openxmlformats-officedocument.wordprocessingml.document')}
    for name in item['attachments']:
        file=directory/name; main,sub=types[file.suffix]
        m.add_attachment(file.read_bytes(),maintype=main,subtype=sub,filename=name)
    m.set_boundary('akten-nachtrag-'+ident)
    (directory/item['file']).write_bytes(m.as_bytes())

def readmes(case,directory):
    inventory=[]
    for item in case['docs']:
        inventory += [(item['file'],item['title']),(str(Path(item['file']).with_suffix('.pdf')),item['title']+' als Lesefassung')]
    inventory += [(e['file'],e['title']) for e in case['emails']]
    inventory += [(case['xlsx'],'Rechenblatt mit Eingaben, Formeln und Belegzuordnung'),(case['chat'],'Ergänzender betrieblicher Nachrichtenverlauf')]
    text = ['# 1 Nachtrag zum 8 Oktober 2026','',
        'Der Kernbestand bleibt beim Aktenstand 2. Oktober 2026. Dieser getrennte Nachtrag führt den fiktiven Vorgang bis zum 8. Oktober fort. Neue Erinnerungen und Verhandlungsstände ersetzen keine älteren Unterlagen.','',
        '## 1.1 Kurzer Einstieg','',case['entry'],'',
        'Die ursprünglichen Workshop-Einstiege und Kernunterlagen bleiben nutzbar. Wer mit dem Nachtrag arbeitet, setzt den Bearbeitungsstand auf den 8. Oktober; Fristen aus den Kernunterlagen werden nicht auf ein späteres Workshopdatum verschoben.','',
        '## 1.2 Arbeitsdateien','',
        'Die sechs Briefe und Protokolle liegen jeweils als bearbeitbares DOCX und inhaltsgleiche PDF-Lesefassung vor. Sechs E-Mails enthalten die bezeichneten Anhänge tatsächlich als MIME-Anlagen. N13 enthält offene Eingaben und Formeln; eine Zahl in der Tabelle entscheidet keine streitige Rechtsfrage.','',
        WARNING,'','| Datei | Inhalt |','| --- | --- |']
    text += [f'| [{name}]({name}) | {title} |' for name,title in inventory]
    text += ['','Alle Personen, Kontakte und Vorgänge sind fiktiv. E-Mail-Adressen enden auf `.example`; Kontobezeichnungen sind interne Belegkennungen. Die Unterlagen enthalten unterschiedliche Parteidarstellungen und keine juristische Musterlösung.','',
             '[Zur Fallübersicht und den Gesamtdownloads](../README.md)','']
    (directory/'README.md').write_text('\n'.join(text),encoding='utf-8')
    readme=directory.parent/'README.md'
    old=readme.read_text(encoding='utf-8')
    start='<!-- BEGIN gesellschafterstreit-nachtrag -->'; end='<!-- END gesellschafterstreit-nachtrag -->'
    section='\n'.join([start,'## 1.5 Vertiefung mit Nachtrag vom 8 Oktober 2026','',
       'Der ursprüngliche Aktenstand vom 2. Oktober und die Kernunterlagen bleiben erhalten. Der zusätzliche Bestand führt denselben Fall bis zum 8. Oktober 2026 fort: sechs ausformulierte DOCX jeweils mit PDF-Lesefassung, sechs E-Mails mit echten Anhängen, eine Formelarbeitsmappe und ein ergänzender Chat.','',case['entry'],'',
       f'[Nachtrag und Dateiverzeichnis]({SUBDIR}/README.md)','',
       'Beim ursprünglichen Kurztermin genügen die Kernunterlagen. Für die Vertiefung werden widersprechende Erinnerungen, Zahlungsvorgänge und offene Verhandlungspositionen hinzugezogen. Die beiden Aktenstände sind ausdrücklich zu unterscheiden.','',end])
    if start in old:old=re.sub(re.escape(start)+r'.*?'+re.escape(end),section,old,flags=re.S)
    else:old=old.rstrip()+'\n\n'+section+'\n'
    readme.write_text(old,encoding='utf-8')

def render(item,path,qa,env):
    dest=qa/'word'/path.parent.parent.name/path.stem;dest.mkdir(parents=True,exist_ok=True)
    proc=subprocess.run([sys.executable,str(RENDER),str(path),'--output_dir',str(dest),'--emit_pdf','--dpi','100'],capture_output=True,text=True,env=env,timeout=300)
    (dest/'render.log').write_text(proc.stdout+proc.stderr)
    images=sorted(dest.glob('page-*.png'));pdfs=list(dest.glob('*.pdf'))
    if proc.returncode or not images or len(pdfs)!=1:raise RuntimeError(str(path)+': '+proc.stderr[-1500:])
    reader=PdfReader(pdfs[0]);w=PdfWriter();w.clone_document_from_reader(reader)
    w.add_metadata({'/Title':item['title'],'/Author':item['sender'].splitlines()[0],'/CreationDate':'D:20261008120000Z','/ModDate':'D:20261008120000Z'})
    if '/Metadata' in w._root_object:del w._root_object[NameObject('/Metadata')]
    raw=hashlib.sha256(path.read_bytes()).digest()[:16];w._ID=ArrayObject([ByteStringObject(raw),ByteStringObject(raw)])
    with path.with_suffix('.pdf').open('wb') as out:w.write(out)
    return dict(file=str(path),pages=len(reader.pages),images=[str(p) for p in images])

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--qa-dir',type=Path,default=Path('/tmp/gs-nachtrag-qa'))
    ap.add_argument('--skip-render',action='store_true',help='Nur für erneutes Erzeugen identischer DOCX bei bereits vorhandenen PDFs')
    ap.add_argument('--native-only',action='store_true',help='Erzeugt DOCX/TXT vor der PDF-Stufe')
    ap.add_argument('--case',choices=[c['slug'] for c in CASES],help='Nur einen Nachtrag erzeugen')
    args=ap.parse_args();args.qa_dir.mkdir(parents=True,exist_ok=True)
    cases=[c for c in CASES if not args.case or c['slug']==args.case]
    spec=importlib.util.spec_from_file_location('render_helpers',ROOT/'scripts/render-startup-gruender-werkstatt.py')
    helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
    env=helper.renderer_environment(args.qa_dir,RUNTIME)
    env['PYTHONPATH']=os.environ.get('PYTHONPATH','')
    jobs=[]
    for case in cases:
        directory=ROOT/'testakten'/case['slug']/SUBDIR;directory.mkdir(exist_ok=True)
        for item in case['docs']:
            path=directory/item['file'];word(item,path);jobs.append((item,path))
        (directory/case['chat']).write_text(case['chat_body'],encoding='utf-8')
    if args.native_only:return
    if not args.skip_render:
        with ThreadPoolExecutor(max_workers=2) as pool:
            manifest=list(pool.map(lambda job:render(*job,args.qa_dir,env),jobs))
        manifest_path=args.qa_dir/'word-render.json'
        if args.case and manifest_path.exists():
            manifest=[x for x in json.loads(manifest_path.read_text()) if '/'+args.case+'/' not in x['file']]+manifest
        (args.qa_dir/'word-render.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False))
        print('DOCX gerendert:',len(manifest),'Seiten:',sum(x['pages'] for x in manifest),flush=True)
    for case in cases:
        directory=ROOT/'testakten'/case['slug']/SUBDIR
        for item in case['emails']:message(item,directory)
        readmes(case,directory)
        print(case['slug'], '18 Dokument-/Maildateien und Chat; XLSX separat',flush=True)

if __name__=='__main__':main()
