#!/usr/bin/env python3
"""Rechen-, Beleg- und Verpackungsregressionen der Betriebskostenfälle."""

import csv
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from email import policy
from email.parser import BytesParser

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
SPEC = importlib.util.spec_from_file_location("bk_belege", ROOT / "scripts/build-betriebskosten-schoeneberg-belege.py")
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)


class BetriebskostenTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        BUILDER.ROOT = Path(cls.temp.name)
        cls.cases = [BUILDER.build(case) for case in BUILDER.CASES]

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_monthly_bills_and_distinct_sources(self):
        for case in self.cases:
            bills = case["belege"]
            self.assertEqual(len([b for b in bills if b["category"] == "Hausservice"]), 12)
            self.assertEqual(len([b for b in bills if b["category"] == "Allgemeinstrom"]), 12)
            self.assertEqual(len({b["number"] for b in bills}), len(bills))
            self.assertEqual(len({b["file"] for b in bills}), len(bills))
            self.assertEqual(sum(b["amount"] for b in bills if b["category"] == "Hausservice"), case["hm"] * 12)

    def test_ground_tax_chain_and_areas(self):
        for case in self.cases:
            self.assertEqual(BUILDER.rounded(case["tax_value"] * .00031 * 4.7), case["tax"])
            self.assertEqual(sum(case["areas"]), case["area"])
            self.assertEqual(case["areas"][int(case["unit"]) - 1], 108)
            path = BUILDER.ROOT / "testakten" / case["slug"] / "belege/Grundsteuer_2025.pdf"
            text = "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
            self.assertIn(BUILDER.money(case["tax"]), text)
            self.assertIn("470 Prozent", text)

    def test_twelve_payments_are_preserved_separately_from_statement(self):
        for case in self.cases:
            self.assertEqual(12 * (case["cold"] + case["heat"]), 3600)
            self.assertEqual(BUILDER.rounded(case["draft_total"] - case["draft_credited"]), case["draft_balance"])
        self.assertEqual(self.cases[0]["draft_credited"], 3300)
        self.assertEqual(self.cases[1]["draft_credited"], 3600)

    def test_csv_columns_and_source_links(self):
        for case in self.cases:
            directory = BUILDER.ROOT / "testakten" / case["slug"]
            for path in directory.rglob("*.csv"):
                with path.open(encoding="utf-8") as stream:
                    rows = list(csv.reader(stream, delimiter=";"))
                self.assertGreaterEqual(len(rows), 5)
                self.assertTrue(all(len(row) == len(rows[0]) for row in rows))
            for bill in case["belege"]:
                self.assertTrue((directory / bill["file"]).is_file())
            with (directory / "zahlenwerk/Flaechen_und_Messwerte_2025.csv").open(encoding="utf-8") as stream:
                headers = next(csv.reader(stream, delimiter=";"))
            self.assertIn("Heizkostenverteiler_Einheiten", headers)
            self.assertNotIn("Heizwärme_kWh", headers)

    def test_substantial_readable_single_documents(self):
        for path in BUILDER.ROOT.rglob("*.pdf"):
            text = "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
            self.assertGreater(len(text), 600, path.name)
            self.assertNotIn(chr(167), text)
            for term in ("Musterlösung", "Lösungsmatrix", "Dozentenhinweis", "Platzhalter", "example.org"):
                self.assertNotIn(term, text, path.name)

    def test_rebuild_is_byte_stable(self):
        before = {p.relative_to(BUILDER.ROOT): hashlib.sha256(p.read_bytes()).hexdigest() for p in BUILDER.ROOT.rglob("*.pdf")}
        for case in BUILDER.CASES:
            BUILDER.build(case)
        after = {p.relative_to(BUILDER.ROOT): hashlib.sha256(p.read_bytes()).hexdigest() for p in BUILDER.ROOT.rglob("*.pdf")}
        self.assertEqual(before, after)

    def test_workbook_sources_formulas_and_cached_totals(self):
        for case in self.cases:
            path = ROOT / "testakten" / case["slug"] / "zahlenwerk/Verwaltungsbuchhaltung_2025.xlsx"
            values = load_workbook(path, data_only=True, read_only=True)
            formulas = load_workbook(path, data_only=False, read_only=True)
            try:
                self.assertEqual(values.sheetnames, ["Belegeingang", "Mietkonto2025", "Abrechnung_Stand"])
                rows = list(values["Belegeingang"].iter_rows(min_row=8, max_row=7 + len(case["belege"]), values_only=True))
                self.assertEqual(len(rows), len(case["belege"]))
                for row, bill in zip(rows, case["belege"]):
                    self.assertEqual(row[5], bill["number"])
                    self.assertAlmostEqual(row[3], bill["amount"], places=2)
                    self.assertEqual(row[9], bill["file"])
                for cell, expected in (("E28", case["draft_total"]), ("E29", case["draft_credited"]), ("E30", case["draft_balance"])):
                    self.assertAlmostEqual(values["Abrechnung_Stand"][cell].value, expected, places=2)
                self.assertEqual(values["Mietkonto2025"]["F21"].value, 3600)
                self.assertTrue(formulas["Abrechnung_Stand"]["E30"].value.startswith("="))
                for sheet in values:
                    for row in sheet:
                        self.assertTrue(all(cell.data_type != "e" for cell in row), sheet.title)
            finally:
                values.close()
                formulas.close()

    def test_native_correspondence_and_separate_documents(self):
        for case in self.cases:
            directory = ROOT / "testakten" / case["slug"]
            self.assertEqual(len(list((directory / "aktenstuecke").glob("*.docx"))), 6)
            self.assertFalse(list((directory / "aktenstuecke").glob("*.md")))
            messages = list((directory / "korrespondenz").glob("*.eml"))
            self.assertEqual(len(messages), 10)
            for path in messages:
                message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
                self.assertFalse(message.defects, path.name)
                for header in ("From", "To", "Subject", "Date", "Message-ID", "MIME-Version"):
                    self.assertTrue(message[header], (path.name, header))
                self.assertGreater(len(message.get_body(preferencelist=("plain",)).get_content()), 900)

    def test_plugin_entry_and_scope(self):
        directory = ROOT / "betriebskosten-hausverwaltung"
        manifest = json.loads((directory / ".claude-plugin/plugin.json").read_text())
        self.assertEqual(manifest["name"], directory.name)
        self.assertLessEqual(len(manifest["description"]), 300)
        self.assertEqual(len(list((directory / "skills").glob("*/SKILL.md"))), 10)
        self.assertTrue((directory / "skills/belege-bis-zur-abrechnung/SKILL.md").is_file())
        mini = (directory / f"{directory.name}-schnellstart.md").read_bytes()
        self.assertLess(len(mini), 7500)
        for path in directory.glob("*-*.md"):
            self.assertNotIn("skills/", str(path.parent.relative_to(directory)))


if __name__ == "__main__":
    unittest.main()
