#!/usr/bin/env python3
"""Unabhängige Integritätsprüfungen der Schnittflug-Nachreichungen.

Liest die ausgelieferten Dateien; baut oder verändert keine Artefakte.
Aufruf: python3 scripts/test-schnittflug-erweiterung.py
Alle Repo-Pfade werden relativ zu diesem Skript bestimmt.
"""
from __future__ import annotations

import csv
from datetime import datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
import hashlib
import json
from pathlib import Path
import re
import unittest
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'scripts/data/startup-gruender'
NATIVE_SUFFIXES = {'.eml', '.docx', '.xlsx', '.pdf', '.png', '.txt', '.csv'}
BERLIN_CUTOFF_OFFSET = timezone(timedelta(hours=2))


def read_data(name):
    return json.loads((DATA / name).read_text(encoding='utf-8'))


def filename(row):
    return row if isinstance(row, str) else row.get('file', row.get('filename'))


def message_id(index):
    return f'<schnittflug-2026-{index:02d}@schnittflug.example>'


def cents(value):
    return int((Decimal(str(value)) * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def euro_text(value):
    whole, fraction = divmod(abs(value), 100)
    return ('-' if value < 0 else '') + f'{whole:,}'.replace(',', '.') + f',{fraction:02d}'


def docx_text(path):
    with ZipFile(path) as archive:
        xml = ElementTree.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    return '\n'.join(''.join(p.itertext()) for p in xml.findall('.//w:p', ns))


def document_datetime(date_label):
    """Sichtbarer Dokumentstand; bei Protokollen zählt die späteste Uhrzeit."""
    months = {'Januar': 1, 'Februar': 2, 'März': 3, 'April': 4, 'Mai': 5,
              'Juni': 6, 'Juli': 7, 'August': 8, 'September': 9,
              'Oktober': 10, 'November': 11, 'Dezember': 12}
    dates = [datetime(int(year), months[month], int(day), tzinfo=BERLIN_CUTOFF_OFFSET)
             for day, month, year in re.findall(r'(\d{1,2})\.?\s+(' + '|'.join(months) + r')\s+(\d{4})', date_label)]
    dates.extend(datetime(int(year), int(month), int(day), tzinfo=BERLIN_CUTOFF_OFFSET)
                 for day, month, year in re.findall(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', date_label))
    if not dates:
        raise ValueError(f'Nicht erkannter sichtbarer Dokumentstand: {date_label}')
    times = [(int(h), int(m)) for h, m in re.findall(r'\b(\d{1,2}):(\d{2})\b', date_label)]
    hour, minute = max(times, default=(0, 0))
    return max(dates).replace(hour=hour, minute=minute)


class SchnittflugErweiterung(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.facts = read_data('fallstamm.json')
        cls.case = ROOT / 'testakten' / cls.facts['case_slug']
        cls.baseline = read_data('morgenstand-originale.json')['files']
        cls.prior = read_data('kommunikation.json')
        cls.communication = read_data('kommunikation-ergaenzung.json')
        cls.contracts = read_data('vertraege-ergaenzung.json')
        cls.finance = read_data('finanz-ergaenzung.json')
        cls.old_finance = read_data('finanz.json')
        cls.mail_rows = cls.prior['mails'] + cls.communication['mails']
        cls.people = cls.prior['people'] | cls.communication['people']
        cls.new_documents = cls.contracts['documents'] + cls.communication['documents']
        cls.inventory_groups = {
            'Morgenstand': [r['file'] for r in cls.baseline],
            'Vertragsnachreichung': [r['file'] for r in cls.contracts['documents']],
            'Kommunikationsnachreichung': [r['file'] for key in ('documents', 'mails', 'chats', 'screenshots', 'registers')
                                            for r in cls.communication[key]],
            'Finanznachreichung': [r['filename'] for r in cls.finance['documents']]
                                  + [filename(r) for r in cls.finance['workbooks']],
        }
        cls.cutoff = datetime.fromisoformat(cls.finance['cutoff']).replace(tzinfo=BERLIN_CUTOFF_OFFSET)

    def test_baseline_hashes_and_exact_manifest_inventory(self):
        declared = [name for group in self.inventory_groups.values() for name in group]
        self.assertEqual(len(declared), len(set(declared)), 'Manifestdateien überlappen sich')
        for name in declared:
            self.assertEqual(Path(name).name, name, name)
            self.assertIn(Path(name).suffix, NATIVE_SUFFIXES, name)
        actual = {p.name for p in self.case.iterdir() if p.is_file() and p.suffix.lower() in NATIVE_SUFFIXES}
        self.assertSetEqual(actual, set(declared))
        for row in self.baseline:
            with self.subTest(original=row['file']):
                self.assertEqual(hashlib.sha256((self.case / row['file']).read_bytes()).hexdigest(), row['sha256'])
        expected_mails = {row['file'] for row in self.mail_rows}
        self.assertEqual(len(expected_mails), len(self.mail_rows))
        self.assertSetEqual({name for name in actual if name.endswith('.eml')}, expected_mails)

    def test_mail_headers_threads_body_and_attachment_bytes(self):
        seen_ids = set()
        attachment_count = 0
        for index, row in enumerate(self.mail_rows, 1):
            with self.subTest(mail=row['file']):
                msg = BytesParser(policy=policy.default).parsebytes((self.case / row['file']).read_bytes())
                self.assertEqual(msg.defects, [])
                self.assertEqual(str(msg['Subject']), row['subject'])
                self.assertEqual(str(msg['Message-ID']), message_id(index))
                self.assertNotIn(str(msg['Message-ID']), seen_ids)
                seen_ids.add(str(msg['Message-ID']))
                for header, field in [('From', 'sender'), ('To', 'to')]:
                    expected = self.people.get(row[field], row[field])
                    self.assertEqual(getaddresses(msg.get_all(header, [])), getaddresses([expected]))
                cc = [self.people.get(value, value) for value in row.get('cc', [])]
                self.assertEqual(getaddresses(msg.get_all('Cc', [])), getaddresses(cc))
                for _, address in getaddresses([str(msg[h]) for h in ('From', 'To', 'Cc') if msg[h] is not None]):
                    self.assertTrue(address.endswith('.example'), address)
                sent = parsedate_to_datetime(msg['Date'])
                self.assertEqual(sent, datetime.fromisoformat(row['date']))
                self.assertIsNotNone(sent.utcoffset())
                self.assertLess(sent, self.cutoff + timedelta(days=1))
                ancestors = []
                ancestor = row.get('reply')
                while ancestor:
                    self.assertGreater(ancestor, 0)
                    self.assertLess(ancestor, index)
                    self.assertNotIn(ancestor, ancestors, 'Zyklischer Mailthread')
                    ancestors.append(ancestor)
                    self.assertLessEqual(datetime.fromisoformat(self.mail_rows[ancestor - 1]['date']), sent)
                    ancestor = self.mail_rows[ancestor - 1].get('reply')
                if ancestors:
                    self.assertEqual(str(msg['In-Reply-To']), message_id(ancestors[0]))
                    self.assertEqual(str(msg['References']).split(), [message_id(n) for n in reversed(ancestors)])
                else:
                    self.assertIsNone(msg['In-Reply-To'])
                    self.assertIsNone(msg['References'])
                body = msg.get_body(preferencelist=('plain',))
                self.assertIsNotNone(body)
                self.assertTrue(body.get_content().replace('\r\n', '\n').startswith(row['body']), row['file'])
                parts = list(msg.iter_attachments())
                self.assertEqual([part.get_filename() for part in parts], row['attachments'])
                for part in parts:
                    name = part.get_filename()
                    self.assertEqual(Path(name).name, name)
                    self.assertEqual(part.get_payload(decode=True), (self.case / name).read_bytes(), name)
                    attachment_count += 1
        self.assertEqual(attachment_count, sum(len(row['attachments']) for row in self.mail_rows))

    def test_actual_docx_content_and_attachment_chronology(self):
        document_stands = {}
        for row in self.new_documents:
            with self.subTest(document=row['file']):
                text = docx_text(self.case / row['file'])
                for required in (row['title'], row['author'], row['date'], row['intro']):
                    self.assertIn(required, text)
                for section in row['sections']:
                    self.assertIn(section['title'], text)
                    for paragraph in section['paragraphs']:
                        self.assertIn(paragraph, text)
                document_stands[row['file']] = document_datetime(row['date'])
                self.assertLess(document_stands[row['file']], self.cutoff + timedelta(days=1))
        for row in self.communication['mails']:
            sent = datetime.fromisoformat(row['date'])
            for name in row['attachments']:
                if name in document_stands:
                    self.assertLessEqual(document_stands[name], sent, f'{row["file"]}: Anhang {name} entsteht erst später')
                elif name in {filename(r) for r in self.finance['workbooks']}:
                    self.assertLessEqual(datetime.fromisoformat(self.finance['received_at']), sent, name)

    def test_native_csv_registers_match_rows_and_existing_sources(self):
        for row in self.communication['registers']:
            with self.subTest(register=row['file']):
                with (self.case / row['file']).open(encoding='utf-8-sig', newline='') as stream:
                    rows = list(csv.reader(stream))
                self.assertEqual(rows, [row['columns']] + row['rows'])
                self.assertTrue(all(len(cells) == len(row['columns']) for cells in rows))
                for name in re.findall(r'\b\d{2,3}_[\w.-]+\.(?:docx|xlsx|pdf|eml|csv|txt|png)\b', str(rows)):
                    self.assertTrue((self.case / name).is_file(), name)

    def test_receipt_amounts_in_actual_pdfs_and_no_duplicate_old_costs(self):
        old = self.old_finance['receipts']
        documents = self.finance['documents']
        costs = [r for r in documents if r['kind'] in ('invoice', 'credit')]
        payments = [r for r in documents if r['kind'] == 'payment']
        self.assertEqual(len(costs) + len(payments), len(documents))
        self.assertEqual(len({r['id'] for r in documents}), len(documents))
        total_cost = sum(r['gross_cents'] for r in old + costs)
        total_paid = sum(r['gross_cents'] for r in old if r['paid_date']) + sum(r['payment_cents'] for r in payments)
        self.assertEqual(total_cost, 389689)
        self.assertEqual(total_paid, 326161)
        self.assertEqual(total_cost - total_paid, 63528)
        old_targets = {r['target']: r for r in payments if r['target'] in {'SF-A007', 'SF-A010'}}
        self.assertEqual(set(old_targets), {'SF-A007', 'SF-A010'})
        self.assertEqual(old_targets['SF-A007']['payment_cents'], 20000)
        self.assertEqual(old_targets['SF-A010']['payment_cents'], 21420)
        self.assertFalse(any(r['target'] in old_targets for r in costs))
        for row in documents:
            with self.subTest(receipt=row['filename']):
                text = ' '.join(' '.join(page.extract_text() or '' for page in PdfReader(self.case / row['filename']).pages).split())
                for required in (row['id'], row['number'], row['recipient'], f'Bezug: {row["target"]}'):
                    self.assertIn(required, text)
                self.assertIn(datetime.fromisoformat(row['date']).strftime('%d.%m.%Y'), text)
                if row['kind'] == 'payment':
                    self.assertIn('Privat gezahlter Betrag ' + euro_text(row['payment_cents']) + ' EUR', text)
                    self.assertNotIn('gross_cents', row, 'Zahlung darf keinen weiteren Kostenbetrag darstellen')
                    self.assertLessEqual(datetime.fromisoformat(row['paid_date']).date(), self.cutoff.date())
                else:
                    self.assertEqual(row['net_cents'] + row['vat_cents'], row['gross_cents'])
                    self.assertEqual(int((Decimal(row['net_cents']) * Decimal('0.19')).quantize(Decimal('1'), rounding=ROUND_HALF_UP)), row['vat_cents'])
                    for label, key in [('Nettobetrag', 'net_cents'), ('Umsatzsteuer 19 %', 'vat_cents'),
                                       ('Gutschrift gesamt' if row['kind'] == 'credit' else 'Rechnungsbetrag', 'gross_cents')]:
                        self.assertIn(label + ' ' + euro_text(row[key]) + ' EUR', text)

    def test_actual_expense_ledger_and_bridge(self):
        expense_file, = [filename(r) for r in self.finance['workbooks'] if filename(r).startswith('140_')]
        book = load_workbook(self.case / expense_file, data_only=True)
        try:
            founders = {r['id']: r['name'] for r in self.facts['founders']}
            expected_costs = {r['id']: r for r in self.old_finance['receipts']}
            expected_costs.update({r['id']: r for r in self.finance['documents'] if r['kind'] in ('invoice', 'credit')})
            actual_costs = {row[1].value: row for row in book['Kosten'] if str(row[1].value).startswith('SF-A')}
            self.assertEqual(set(actual_costs), set(expected_costs))
            self.assertEqual(len(actual_costs), sum(str(row[1].value).startswith('SF-A') for row in book['Kosten']))
            for key, record in expected_costs.items():
                row = actual_costs[key]
                target = record.get('target', 'SF-A005' if key == 'SF-A011' else key)
                self.assertEqual(row[2].value, target, key)
                self.assertEqual(row[3].value, founders[record['founder_id']], key)
                self.assertEqual([cents(row[c].value) for c in (5, 6, 7)], [record[k] for k in ('net_cents', 'vat_cents', 'gross_cents')], key)
                self.assertEqual(row[9].value, record['filename'], key)
            expected_payments = {r['id'] + '-Z': (r['id'] if r['id'] != 'SF-A011' else 'SF-A005', r['gross_cents'], r['filename'])
                                 for r in self.old_finance['receipts'] if r['paid_date']}
            expected_payments.update({r['id']: (r['target'], r['payment_cents'], r['filename'])
                                      for r in self.finance['documents'] if r['kind'] == 'payment'})
            actual_payments = {row[1].value: (row[2].value, cents(row[5].value), row[8].value)
                               for row in book['Zahlungen'] if str(row[1].value).startswith('SF-A')}
            self.assertEqual(actual_payments, expected_payments)
            self.assertEqual(len(actual_payments), sum(str(row[1].value).startswith('SF-A') for row in book['Zahlungen']))
            a = book['Abgleich']
            for line, expected in [(7, (310923, 239523, 71400)), (8, (78766, 86638, -7872)),
                                   (9, (389689, 326161, 63528)), (33, (389689, 326161, 63528))]:
                self.assertEqual(tuple(cents(a[f'{col}{line}'].value) for col in 'DEF'), expected)
            self.assertEqual(cents(a['J33'].value), 0)
            balances = {a[f'B{r}'].value: (cents(a[f'D{r}'].value), cents(a[f'E{r}'].value), cents(a[f'F{r}'].value)) for r in range(16, 32)}
            self.assertEqual(balances['SF-A007'], (49980, 20000, 29980))
            self.assertEqual(balances['SF-A010'], (21420, 21420, 0))
            self.assertEqual(balances['SF-A014'], (49980, 25000, 24980))
            self.assertEqual(balances['SF-A016'], (17850, 17850, 0))
            self.assertEqual(balances['SF-A021'], (8568, 0, 8568))
            self.assertTrue(all(a[f'{col}{r}'].value == 'Offen' for col in 'GHI' for r in range(16, 32)))
        finally:
            book.close()

    def test_financing_conditions_and_offer_are_not_receipts(self):
        funding_file, = [filename(r) for r in self.finance['workbooks'] if filename(r).startswith('141_')]
        book = load_workbook(self.case / funding_file, data_only=True)
        try:
            f = book['Finanzierung']
            self.assertEqual((f['C7'].value, f['C8'].value), (25000, 600000))
            self.assertEqual((f['E7'].value, f['E8'].value), (0, 0))
            self.assertEqual((f['E25'].value, f['E26'].value), (3, 4))
            self.assertTrue(all(f[f'E{r}'].value == 'Offen' for r in range(16, 23)))
            due = book['Fälligkeiten']
            self.assertEqual(cents(due['G12'].value), 63528)
            self.assertEqual([cents(due[f'E{r}'].value) for r in (19, 20, 21, 23)], [856800, 856800, 428400, 2142000])
            self.assertTrue(all(due[f'F{r}'].value == 'Angebotsannahme' for r in range(19, 22)))
            uses = book['Mittelverwendung']
            self.assertEqual([cents(uses[f'C{r}'].value) for r in (6, 7, 8, 9, 10)], [60000000, 56500000, 3500000, 326161, 3173839])
            self.assertEqual(cents(uses['G28'].value), 636500)
        finally:
            book.close()

    def test_every_delivered_workbook_has_cached_formula_values(self):
        for path in sorted(self.case.glob('*.xlsx')):
            with self.subTest(workbook=path.name):
                formulas = load_workbook(path, data_only=False)
                values = load_workbook(path, data_only=True)
                try:
                    count = 0
                    self.assertEqual(formulas.sheetnames, values.sheetnames)
                    for sheet in formulas:
                        for row in sheet:
                            for cell in row:
                                cached = values[sheet.title][cell.coordinate]
                                label = f'{path.name}/{sheet.title}/{cell.coordinate}'
                                self.assertNotEqual(cached.data_type, 'e', label)
                                if cell.data_type == 'f':
                                    count += 1
                                    self.assertIsNotNone(cached.value, label)
                                    self.assertFalse(isinstance(cached.value, str) and cached.value.startswith('='), label)
                    self.assertGreater(count, 0, path.name)
                finally:
                    formulas.close()
                    values.close()


if __name__ == '__main__':
    unittest.main()
