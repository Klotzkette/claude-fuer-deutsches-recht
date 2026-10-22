#!/usr/bin/env python3
"""Prüft Quellentreue, Anlagenaufteilung, Originalformate und Gleichlauf beider Fassungen."""

import csv
from datetime import date
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime, getaddresses
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import sys
import zipfile
from unittest.mock import patch

from docx import Document
from PIL import Image
from pypdf import PdfReader
from openpyxl import load_workbook
from testakte_disclaimer import NOTICE_BYTES, pdf_content_errors

ROOT = Path(__file__).resolve().parents[1]
PLUGIN = ROOT / "arbeitsrecht-kuendigungsschutzklagewerkstatt"
spec = importlib.util.spec_from_file_location("seewerk_build", ROOT / "scripts/build-seewerk-akten.py")
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


def require(condition, message):
    if not condition:
        raise ValueError(message)


def normalized(value):
    return re.sub(r"\s+", " ", value).strip()


def text_of(path):
    if path.suffix == ".docx":
        document = Document(path)
        return "\n".join([*(paragraph.text for paragraph in document.paragraphs),
                          *(cell.text for table in document.tables for row in table.rows for cell in row.cells)])
    if path.suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        return message.get_body(preferencelist=("plain",)).get_content()
    if path.suffix == ".xlsx":
        book = load_workbook(path, data_only=True)
        return "\n".join(str(cell.value) for row in book.active for cell in row if cell.value is not None)
    if path.suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if path.suffix == ".png":
        return ""
    return path.read_text(encoding="utf-8-sig")


def pdf_texts(data):
    require(not pdf_content_errors(data), "Warnhinweis oder beschädigtes PDF")
    return [normalized(page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages]


def check():
    source = json.loads(builder.DATA.read_text(encoding="utf-8"))
    core = builder.core_documents(source["blocks"])
    reports = json.loads((PLUGIN / "testakte/pruefung/seewerk-bestand.json").read_text(encoding="utf-8"))
    require([item["slug"] for item in reports] == builder.SLUGS, "Andere Fallfassungen")
    require(re.fullmatch(r"[a-f0-9]{64}", source["source_sha256"]), "Quellbezug fehlt")
    quality = builder.module("validate-testakten-dokumentqualitaet")
    exceptions = quality.reviewed_source_documents()
    require(len(exceptions) == 16, "Einzeln begrenzte Quellbeleg-Ausnahmen")
    for relative, rule in exceptions.items():
        path = ROOT / relative
        require(quality.reviewed_source_rule(path) == rule, "Geprüfter Beleg verändert")
        with patch.object(quality.hashlib, "sha256") as digest:
            digest.return_value.hexdigest.return_value = "0"*64
            require(quality.reviewed_source_rule(path) == {}, "Ausnahme gilt für veränderte Datei")
            require(quality.minimum_document_text(path) == 600, "Mindestlänge pauschal abgeschwächt")
    for index, report in enumerate(reports):
        directory = ROOT / "testakten" / report["slug"]
        readme = (directory / "README.md").read_text(encoding="utf-8")
        for target in ("../README.md", "../../README.md", "../../ASSET_INDEX.md"):
            require("](" + target + ")" in readme, "Aktennavigation fehlt: " + target)
        files = report["files"]
        require(report["documents"] == len(files) == (23 if index == 0 else 44), "Dokumentenzahl")
        require(len({item["file"] for item in files}) == len(files), "Dateidoppelung")
        for item in core:
            path = directory / item["filename"]
            actual = normalized(text_of(path))
            for block in item["blocks"]:
                parts = [block["text"]] if "text" in block else [cell for row in block["rows"] for cell in row]
                for part in parts:
                    require(normalized(part) in actual, f"Kerninhalt fehlt oder verändert: {path.name}: {part[:90]}")
        all_text = []
        for entry in files:
            path = directory / entry["file"]
            data = path.read_bytes()
            require(hashlib.sha256(data).hexdigest() == entry["sha256"], f"Original-Prüfsumme: {path.name}")
            require(re.fullmatch(r"[A-Za-z0-9_]+\.(docx|xlsx|csv|eml|txt|pdf|png)", path.name), "Ungeeigneter Dateiname")
            text = text_of(path)
            all_text.append(text)
            if path.suffix in {".docx", ".eml", ".pdf", ".txt"}:
                inspected = quality.docx_text(path) if path.suffix == ".docx" else text
                findings = quality.export_meta_errors(inspected,path) + quality.language_prose_errors(inspected,path)
                if path.suffix == ".eml":
                    findings += quality.eml_quality_errors(path)
                require(not findings, f"Dokumentqualität {path.name}: {findings}")
                require(len(inspected.strip()) >= quality.minimum_document_text(path), f"Kurzunterlage ohne Quellbelegfreigabe: {path.name}")
                if path.suffix == ".pdf":
                    require(quality.pdf_is_a4(path), "PDF nicht auf A4")
            require(not re.search(r"§|Lösungsmatrix|Musterlösung|Formathinweis|Lorem|\bTODO\b|Eberle|Pause machen", text), f"Spoiler oder Fremdtext: {path.name}")
            if path.suffix == ".docx":
                doc = Document(path)
                require(doc.styles["Normal"].font.name == "Times New Roman" and doc.styles["Normal"].font.size.pt == 11, "Schriftstandard")
                require(doc.core_properties.author == "Klotzkette", "Autor")
                require(doc.core_properties.created.date() <= date(2026,11,23), "Spätere Unterlage")
            elif path.suffix == ".eml":
                mail = BytesParser(policy=policy.default).parsebytes(data)
                require(not mail.defects and all(mail[value] for value in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version")), "Mailheader")
                require(parsedate_to_datetime(mail["Date"]).date() <= date(2026,11,23), "Spätere E-Mail")
                require(all(address.endswith(".example") for _, address in getaddresses([mail["From"],mail["To"]])), "Kontaktadresse")
                for attachment in mail.iter_attachments():
                    require(attachment.get_payload(decode=True) == (directory / attachment.get_filename()).read_bytes(), "Veraltete MIME-Anlage")
            elif path.suffix == ".csv":
                rows = list(csv.reader(io.StringIO(text),delimiter=";"))
                require(len(rows) >= 9 and len({len(row) for row in rows}) == 1, "CSV-Inhalt")
            elif path.suffix == ".png":
                image = Image.open(path)
                require(image.size == (1480,1120) and len(image.getcolors(image.width*image.height)) > 100, "Leerer Screenshot")
        combined_text = "\n".join(all_text)
        for needle in ("6.696", "06:12", "06:24", "07:20", "07:48", "am Dienstag", "am Donnerstag", "30 Tage", "72 Stunden", "i. A.", "30.04.2027"):
            require(needle in combined_text, f"Vorgegebener Aktenumstand verloren: {needle}")
        prefix = PLUGIN / "testakte/downloads" / ("testakte-" + report["slug"])
        pages = []
        with zipfile.ZipFile(str(prefix)+".zip") as archive:
            require(archive.namelist()[0] == "README.txt" and archive.read("README.txt") == NOTICE_BYTES, "Originalhinweis")
            require(set(archive.namelist()) == {item["file"] for item in files} | {"README.txt"}, "Originalarchivumfang")
            require(all("/" not in name and "\\" not in name and not name.endswith(".md") for name in archive.namelist()), "Archiv nicht flach")
            for item in files:
                require(archive.read(item["file"]) == (directory / item["file"]).read_bytes(), "Originalarchiv veraltet")
        with zipfile.ZipFile(str(prefix)+"-einzelpdfs.zip") as archive:
            require(archive.namelist()[0] == "README.txt" and archive.read("README.txt") == NOTICE_BYTES, "PDF-Hinweis")
            require(set(archive.namelist()) == {Path(item["file"]).stem+".pdf" for item in files} | {"README.txt"}, "Ein Dokument je PDF fehlt")
            for item in files:
                pdf = archive.read(Path(item["file"]).stem+".pdf")
                require(hashlib.sha256(pdf).hexdigest() == item["pdf_sha256"], "PDF-Prüfsumme")
                texts = pdf_texts(pdf)
                require(len(texts) == item["pages"], "Seitenumfang")
                if not item["file"].endswith(".png"):
                    require(all(len(value) > 20 for value in texts), "Leere Textseite")
                pages += texts
            with zipfile.ZipFile(str(prefix)+"-anlagen-k1-k14.zip") as annexes:
                require(annexes.namelist() == ["README.txt",*(f"K{number:02d}.pdf" for number in range(1,15))], "Genau vierzehn Anlagen erforderlich")
                require(annexes.read("README.txt") == NOTICE_BYTES, "Anlagenhinweis")
                for number in range(1,15):
                    expected = [text for item in files if item["annex"] == number
                                for text in pdf_texts(archive.read(Path(item["file"]).stem+".pdf"))]
                    labelled = pdf_texts(annexes.read(f"K{number:02d}.pdf"))
                    require(all(value.endswith(f"Anlage K{number}") for value in labelled), "Anlagenkennzeichnung")
                    require([re.sub(r"\s*Anlage K\d+$", "", value) for value in labelled] == expected, "Anlagenzuordnung")
        total = directory / "gesamt-pdf" / (report["slug"]+"_gesamt.pdf")
        require(pdf_texts(total.read_bytes()) == pages and len(pages) == report["pages"], "Gesamtband nicht synchron")
        require((PLUGIN / "testakte/downloads" / total.name).read_bytes() == total.read_bytes(), "PDF-Asset veraltet")
    require(reports[1]["pages"] > reports[0]["pages"]*2, "Ausführliche Fassung ohne substantiellen Ausbau")
    print("Seewerk OK: zwei Fassungen, 67 Originale, vollständiger Kerntext, 28 logische Anlagen, flache Archive und identische PDF-Seiten.")


if __name__ == "__main__":
    check()
