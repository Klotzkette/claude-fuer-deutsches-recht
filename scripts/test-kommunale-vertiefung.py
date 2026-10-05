#!/usr/bin/env python3
"""Inhalts- und Exportregression der kommunalen Fallvertiefung."""
from collections import Counter
from datetime import datetime
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime, getaddresses
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import re
import unittest
import zipfile
from docx import Document
from pypdf import PdfReader
from kommunale_vertiefung_daten import cases, FOLDER, native_names
from kommunale_schriftverkehr_daten import cases as old_cases
from testakte_einzelpdf_common import document_arcname_pairs
from testakte_zip_common import working_dump_archive_pairs
from testakte_disclaimer import NOTICE_BYTES, NOTICE_DE, NOTICE_EN

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('vertiefung', ROOT/'scripts/build-kommunale-vertiefung.py')
b = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b)
norm = lambda s: re.sub(r'\s+', '', s).replace('\u00ad', '')

def pdf_text(reader):
    result = []
    for number, page in enumerate(reader.pages, 1):
        lines = (page.extract_text() or '').strip().splitlines()
        if lines and lines[0].strip() == str(number): lines.pop(0)
        if lines and lines[-1].strip() == str(number): lines.pop()
        result.append('\n'.join(lines))
    return '\n'.join(result)

class Vertiefung(unittest.TestCase):
    def test_all_220_previous_originals_are_unchanged(self):
        frozen = json.loads((ROOT/'quality/kommunale-haftpflicht/vertiefung-altbestand.json').read_text())
        self.assertEqual(frozen['baseline'], 'v445.31.3')
        self.assertEqual(len(frozen['sources']), 220)
        self.assertEqual(len({x['path'] for x in frozen['sources']}), 220)
        for record in frozen['sources']:
            self.assertEqual(hashlib.sha256((ROOT/record['path']).read_bytes()).hexdigest(), record['sha256'], record['path'])

    def test_sixteen_cases_and_precise_new_inventory(self):
        selected = cases()
        self.assertEqual(len(selected), 16)
        self.assertEqual({c['slug'] for c in selected if c['is_new']}, {'akha-wuerzburg-kitaplatz', 'akha-wuerzburg-feuerwehreinsatz', 'akha-wuerzburg-kreisstrassenbaum'})
        for case in selected:
            directory = ROOT/'testakten'/case['slug']/FOLDER
            self.assertEqual({p.name for p in directory.iterdir() if p.is_file()}, set(native_names(case)), case['slug'])
            self.assertGreaterEqual(sum(d['kind'] == 'letter' for d in case['documents']), 2)
            self.assertGreaterEqual(sum(d['kind'] == 'email' for d in case['documents']), 2)
            claims = [directory.parent/'schriftverkehr-und-klage'/d['file'] for old in old_cases() if old['slug'] == case['slug'] for d in old['documents'] if d['kind'] == 'claim'] + [directory/d['file'] for d in case['documents'] if d['kind'] == 'claim']
            self.assertEqual(len(claims), 1, case['slug'])
            self.assertTrue(claims[0].is_file())

    def test_entire_word_text_and_letter_pdf_match(self):
        for case in cases():
            for item in case['documents']:
                if item['kind'] == 'email': continue
                path = ROOT/'testakten'/case['slug']/FOLDER/item['file']
                doc = Document(path)
                text = '\n'.join(p.text for p in doc.paragraphs)
                self.assertEqual(doc.styles['Normal'].font.name, 'Times New Roman')
                self.assertEqual(doc.styles['Normal'].font.size.pt, 11)
                self.assertEqual(sum(p.style.name == 'Title' for p in doc.paragraphs), 1)
                for paragraph in re.split(r'\n\s*\n', b.body(case, item)):
                    self.assertIn(norm(paragraph.removeprefix('## ')), norm(text), (case['slug'], item['file']))
                if item['kind'] == 'letter':
                    pdf = pdf_text(PdfReader(path.with_suffix('.pdf')))
                    for paragraph in doc.paragraphs:
                        self.assertIn(norm(paragraph.text), norm(pdf), (case['slug'], item['file'], paragraph.text[:80]))
                    self.assertNotIn(NOTICE_DE, pdf)
                    self.assertNotIn(NOTICE_EN, pdf)

    def test_existing_letters_have_matching_editable_versions(self):
        for case in old_cases():
            for item in case['documents']:
                if item['kind'] != 'letter': continue
                path = ROOT/'testakten'/case['slug']/'briefvorlagen'/item['file']
                text = '\n'.join(p.text for p in Document(path).paragraphs)
                old_pdf = ROOT/'testakten'/case['slug']/'schriftverkehr-und-klage'/str(Path(item['file']).with_suffix('.pdf'))
                pdf = pdf_text(PdfReader(old_pdf))
                for paragraph in item['body'].split('\n\n'):
                    p = norm(paragraph.removeprefix('## '))
                    self.assertIn(p, norm(text), str(path))
                    self.assertIn(p, norm(pdf), str(old_pdf))

    def test_genuine_mime_dates_recipients_and_attachments(self):
        ids = set()
        for case in cases():
            root = ROOT/'testakten'/case['slug']
            for item in case['documents']:
                if item['kind'] != 'email': continue
                msg = BytesParser(policy=policy.default).parsebytes((root/FOLDER/item['file']).read_bytes())
                self.assertFalse(msg.defects)
                if msg.is_multipart():
                    self.assertLessEqual(len(msg.get_boundary()), 70)
                for h in ('From', 'To', 'Subject', 'Date', 'Message-ID', 'MIME-Version'):
                    self.assertTrue(msg[h])
                    self.assertFalse(msg[h].defects)
                self.assertNotIn(str(msg['Message-ID']), ids)
                ids.add(str(msg['Message-ID']))
                self.assertEqual(parsedate_to_datetime(msg['Date']), datetime.fromisoformat(item['date']))
                self.assertLessEqual(parsedate_to_datetime(msg['Date']).date(), datetime(2026, 10, 5).date())
                for h, key in [('From', 'sender'), ('To', 'recipient')]:
                    self.assertEqual([(x.display_name, x.addr_spec) for x in msg[h].addresses], getaddresses([item[key]]))
                self.assertEqual(msg.get_body(preferencelist=('plain',)).get_content().replace('\r\n', '\n').strip(), item['body'].strip())
                parts = list(msg.iter_attachments())
                expected = {Path(n).name: (root/n).read_bytes() for n in case.get('attachments', {}).get(item['file'], [])}
                self.assertEqual(len(parts), len(expected))
                self.assertEqual({p.get_filename(): p.get_payload(decode=True) for p in parts}, expected)

    def test_claim_exhibits_are_used_and_resolve_to_actual_documents(self):
        for case in cases():
            claims = [d for d in case['documents'] if d['kind'] == 'claim']
            if not claims: continue
            self.assertEqual(len(claims), 1)
            item = claims[0]
            labels = {x['label'] for x in case['exhibits']}
            self.assertEqual(labels, {f'K{i+1}' for i in range(len(labels))})
            self.assertEqual({'K'+n for n in re.findall(r'\bK\s*(\d+)\b', item['body'])}, labels, case['slug'])
            path = ROOT/'testakten'/case['slug']/FOLDER/item['file']
            text = '\n'.join(p.text for p in Document(path).paragraphs)
            self.assertIn('Anlagenverzeichnis', text)
            self.assertRegex(text, r'(?i)nicht eingereicht')
            self.assertIn('BGH', text)
            self.assertIn('Zinsen', text)
            for exhibit in case['exhibits']:
                self.assertTrue(b.source(path.parent.parent, exhibit['source']).is_file())
                self.assertIn(exhibit['source'], text)

    def test_complete_combined_pdf_navigation(self):
        for case in cases():
            root = ROOT/'testakten'/case['slug']
            pdf = PdfReader(root/'gesamt-pdf'/f"{case['slug']}_gesamt.pdf")
            pairs = document_arcname_pairs(root)
            stems = Counter(p.relative_to(root).with_suffix('').as_posix() for p, arc in pairs)
            expected = {p.relative_to(root).as_posix() if stems[p.relative_to(root).with_suffix('').as_posix()] > 1 else p.relative_to(root).with_suffix('').as_posix() for p, arc in pairs}
            outlines = [item for item in pdf.outline if not isinstance(item, list)]
            self.assertEqual(len(outlines), len(expected))
            self.assertEqual({item.title for item in outlines}, expected)
            text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
            self.assertNotIn(NOTICE_DE, text)
            self.assertNotIn(NOTICE_EN, text)

    def test_readme_page_counts_match_final_combined_pdf(self):
        for case in cases():
            root = ROOT/'testakten'/case['slug']
            pages = len(PdfReader(root/'gesamt-pdf'/f"{case['slug']}_gesamt.pdf").pages)
            counts = re.findall(r'Das(?: aktualisierte)? Gesamt-PDF (?:umfasst|enthält) (\d+) Seiten\.', (root/'README.md').read_text())
            self.assertTrue(counts, case['slug'])
            self.assertEqual({int(x) for x in counts}, {pages}, case['slug'])

    @unittest.skipUnless(os.environ.get('KOMMUNALE_VERTIEFUNG_ZIPS'), 'ZIP-Verzeichnis separat bereitstellen')
    def test_archive_inventory_bytes_and_full_converted_text(self):
        packages = Path(os.environ['KOMMUNALE_VERTIEFUNG_ZIPS'])
        for case in cases():
            root = ROOT/'testakten'/case['slug']
            for suffix, pairs in [('', working_dump_archive_pairs(root, include_gesamt_pdf=True)), ('-einzelpdfs', document_arcname_pairs(root))]:
                with zipfile.ZipFile(packages/f"testakte-{case['slug']}{suffix}.zip") as archive:
                    expected = {arc: p for p, arc in pairs}
                    self.assertEqual(set(archive.namelist()), set(expected)|{'README.txt'})
                    self.assertEqual(len(archive.namelist()), len(expected)+1)
                    self.assertTrue(archive.read('README.txt').startswith(NOTICE_BYTES))
                    self.assertTrue(all('/' not in n for n in archive.namelist()))
                    self.assertIsNone(archive.testzip())
                    for arc, path in expected.items():
                        if not suffix or path.suffix == '.pdf': self.assertEqual(archive.read(arc), path.read_bytes())
                    if suffix:
                        reverse = {str(p.relative_to(root)): arc for arc,p in expected.items()}
                        for item in case['documents']:
                            text = pdf_text(PdfReader(io.BytesIO(archive.read(reverse[FOLDER+'/'+item['file']]))))
                            for paragraph in re.split(r'\n\s*\n', b.body(case,item)):
                                self.assertIn(norm(paragraph.removeprefix('## ')), norm(text), (case['slug'], item['file']))

if __name__ == '__main__': unittest.main(verbosity=2)
