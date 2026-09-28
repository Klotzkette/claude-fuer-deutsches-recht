#!/usr/bin/env python3
"""Erstellt ausschließlich die native Göttinger Enteignungsakte."""

from __future__ import annotations

import argparse
import csv
import io
import json
from datetime import datetime, timezone
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from akten_build_runtime import serif_font_path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "scripts/fixtures/enteignung-artikel-14/case.json"
CASE = ROOT / "testakten/enteignung-verkehrsflaeche-goettingen"
FONT_DIR = None


def load_case(path=FIXTURE):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    documents = data["documents"]
    names = [doc["file"] for doc in documents]
    if len(documents) != 11 or len(set(names)) != len(names):
        raise ValueError("Erwartet werden elf eindeutig benannte Originalunterlagen")
    for doc in documents:
        name = Path(doc["file"])
        if name.name != doc["file"] or name.suffix != "." + doc["kind"]:
            raise ValueError("Ungültiger Dokumentpfad oder Dateityp")
        for attachment in doc.get("attachments", []):
            if attachment not in names or not attachment.endswith(".pdf"):
                raise ValueError("Unbekannter PDF-Anhang")
    return data


def register_fonts(font_dir):
    for bold, (label, filename) in enumerate((("CaseRoman", "Times New Roman.ttf"), ("CaseBold", "Times New Roman Bold.ttf"))):
        path = Path(font_dir) / filename if font_dir else serif_font_path(bool(bold))
        if not path.is_file():
            raise FileNotFoundError(f"Schrift fehlt: {path}; --font-dir mit Times New Roman angeben")
        pdfmetrics.registerFont(TTFont(label, str(path)))


def diagram():
    drawing = Drawing(450, 174)
    drawing.add(Rect(5, 6, 435, 153, fillColor=colors.white, strokeColor=colors.black))
    drawing.add(Rect(45, 18, 76, 105, fillColor=colors.HexColor("#ecd3d3")))
    drawing.add(Rect(121, 18, 62, 105, fillColor=colors.HexColor("#d6e5ee")))
    drawing.add(Rect(220, 40, 170, 72, fillColor=colors.HexColor("#eeeeee")))
    drawing.add(Line(5, 138, 440, 138, strokeColor=colors.HexColor("#2879a6")))
    for x, y, label in ((53, 70, "A 620 m²"), (125, 54, "B 180 m²"), (244, 72, "Werkstatthalle"), (204, 143, "Graben / Nordkorridor N"), (194, 24, "Resthof"), (7, 163, "Norden")):
        drawing.add(String(x, y, label, fontName="CaseRoman", fontSize=11))
    return drawing


def pdf_document(doc, path):
    body = ParagraphStyle("body", fontName="CaseRoman", fontSize=11, leading=14.4,
                          spaceAfter=8, alignment=TA_JUSTIFY)
    meta = ParagraphStyle("meta", parent=body, alignment=0, fontSize=10, leading=12, spaceAfter=4)
    title = ParagraphStyle("title", parent=body, fontName="CaseBold", fontSize=15, leading=18, alignment=0, spaceAfter=12)
    heading = ParagraphStyle("heading", parent=body, fontName="CaseBold", fontSize=12, leading=15, alignment=0, spaceAfter=9)
    story = []
    for index, page in enumerate(doc["pages"]):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(escape(doc["title"]), title))
        if index == 0:
            for line in (doc["author"], "An: " + doc["recipient"], doc["date"] + " | " + doc["reference"]):
                story.append(Paragraph(escape(line), meta))
            story.append(Spacer(1, 9))
        story.append(Paragraph(escape(page["heading"]), heading))
        story.extend(Paragraph(escape(text), body) for text in page["paragraphs"])
        if table := page.get("table"):
            rows = [[Paragraph(escape(str(value)), meta) for value in row] for row in [table["headers"], *table["rows"]]]
            widths = [108, 118, 94, 165] if len(table["headers"]) == 4 else [185, 160, 140]
            native = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
            native.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b5b5b5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(native)
        if page.get("diagram"):
            story.append(Spacer(1, 9))
            story.append(diagram())

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("CaseRoman", 9)
        canvas.drawString(55, 35, doc["reference"])
        canvas.drawRightString(A4[0] - 55, 35, f"Seite {document.page}")
        canvas.restoreState()

    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=55, rightMargin=55,
                      topMargin=45, bottomMargin=50, title=doc["title"], author=doc["author"],
                      invariant=1).build(story, onFirstPage=footer, onLaterPages=footer)


def docx_document(doc, path):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Footer"):
        style = document.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attr), None)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attr), "Times New Roman")
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    document.styles["Title"].font.size = Pt(15)
    document.styles["Title"].font.bold = True
    document.styles["Heading 1"].font.bold = True
    document.styles["Normal"].paragraph_format.line_spacing = 1.1
    document.core_properties.author = doc["author"]
    document.core_properties.title = doc["title"]
    document.core_properties.last_modified_by = doc["author"]
    document.core_properties.created = document.core_properties.modified = datetime(2026, 8, 18, tzinfo=timezone.utc)
    document.core_properties.comments = ""
    for index, page in enumerate(doc["pages"]):
        if index:
            document.add_page_break()
        else:
            document.add_paragraph(doc["title"], "Title")
            for line in (doc["author"], "An: " + doc["recipient"], doc["date"] + " | " + doc["reference"]):
                document.add_paragraph(line)
        document.add_paragraph(page["heading"], "Heading 1")
        for text in page["paragraphs"]:
            document.add_paragraph(text)
    footer = section.footer.paragraphs[0]
    footer.add_run(doc["reference"] + " | Seite ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    memory = io.BytesIO()
    document.save(memory)
    # Feste ZIP-Zeiten halten Neuberechnungen und eingebettete Originale vergleichbar.
    with ZipFile(memory) as source, ZipFile(path, "w", compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, (2026, 8, 18, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))


def build(output=CASE, font_dir=FONT_DIR):
    data = load_case()
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    register_fonts(font_dir)
    for doc in data["documents"]:
        path = output / doc["file"]
        if doc["kind"] == "pdf":
            pdf_document(doc, path)
        elif doc["kind"] == "docx":
            docx_document(doc, path)
        elif doc["kind"] == "txt":
            lines = [doc["title"], doc["author"], "An: " + doc["recipient"], doc["date"] + " | " + doc["reference"], *doc["paragraphs"]]
            path.write_text("\n\n".join(lines) + "\n", encoding="utf-8")
        elif doc["kind"] == "csv":
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream, delimiter=";", lineterminator="\r\n")
                writer.writerow(doc["headers"])
                writer.writerows(doc["rows"])
    for doc in data["documents"]:
        if doc["kind"] != "eml":
            continue
        message = EmailMessage(policy=SMTP)
        for header, value in (("From", f'{doc["author"]} <{doc["from"]}>'),
                              ("To", f'{doc["recipient"]} <{doc["to"]}>'),
                              ("Subject", doc["title"]),
                              ("Message-ID", f'<weidenrain-{doc["id"]}-{doc["date"][:10]}@post.example>')):
            message[header] = value
        message["Date"] = datetime.fromisoformat(doc["date"])
        message.set_content("\n\n".join(doc["paragraphs"]) + "\n", charset="utf-8")
        for name in doc.get("attachments", []):
            message.add_attachment((output / name).read_bytes(), maintype="application", subtype="pdf", filename=name)
        if message.is_multipart():
            message.set_boundary(f'weidenrain-{doc["id"]}-originalanlagen')
        (output / doc["file"]).write_bytes(message.as_bytes())
    return [output / doc["file"] for doc in data["documents"]]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=CASE)
    parser.add_argument("--font-dir", type=Path, default=FONT_DIR)
    args = parser.parse_args()
    for path in build(args.output_dir, args.font_dir):
        print(path.relative_to(ROOT) if path.is_relative_to(ROOT) else path)


if __name__ == "__main__":
    main()
