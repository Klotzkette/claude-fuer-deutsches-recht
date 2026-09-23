#!/usr/bin/env python3
"""Fallbegrenzte Quellen-, Rechen- und Exportkontrolle ohne zentrale Änderungen."""

import argparse
from collections import Counter
import csv
from datetime import datetime, timedelta
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from pypdf import PdfReader

from akten_build_runtime import screen_font
from testakte_file_filter import include_in_working_dump
from testakte_office_pdf import render_office_batch, uncached_formula_cells

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / 'testakten/wirtschaftsanwalt-gesellschafterkonflikt-handwerk-hannover'
QA = Path(os.environ['HANNOVER_AKTEN_QA']) if os.environ.get('HANNOVER_AKTEN_QA') else Path(tempfile.mkdtemp(prefix='wirtschaftsanwalt-hannover-test-'))
TOTAL = CASE / 'gesamt-pdf' / (CASE.name + '_gesamt.pdf')


def module(name, filename):
    spec = importlib.util.spec_from_file_location(name, ROOT / 'scripts' / filename)
    loaded = importlib.util.module_from_spec(spec)
    sys.modules[name] = loaded
    spec.loader.exec_module(loaded)
    return loaded


QUALITY = module('hannover_quality', 'validate-testakten-dokumentqualitaet.py')
CSV = module('hannover_csv', 'validate-testakten-csv.py')
README = module('hannover_readme', 'validate-testakten-readme-downloads.py')
SOLUTIONS = module('hannover_solutions', 'validate-testakten-ohne-loesungen.py')


def sources():
    return sorted(p for p in CASE.iterdir() if include_in_working_dump(p, CASE))


def money(value):
    return Decimal(value.replace('.', '').replace(',', '.'))


def rows(prefix):
    with next(CASE.glob(prefix + '*.csv')).open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.DictReader(handle, delimiter=';'))


def recalculation_copy(source, changed, changes):
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(changed, 'w', zipfile.ZIP_DEFLATED) as copy:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename.startswith('xl/worksheets/sheet') and entry.filename.endswith('.xml'):
                tree = ET.fromstring(data)
                for address, value in changes.get(entry.filename, {}).items():
                    tree.find(f'.//{ns}c[@r="{address}"]/{ns}v').text = str(value)
                # Die XML-Eingabeänderung muss abhängige Caches in der Prüfkopie invalidieren.
                for cell in tree.iter(f'{ns}c'):
                    if cell.find(f'{ns}f') is not None:
                        cached = cell.find(f'{ns}v')
                        if cached is not None:
                            cell.remove(cached)
                data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            copy.writestr(entry, data)


def render_sources():
    office = [p for p in sources() if p.suffix in {'.docx', '.xlsx'}]
    office_pdfs = QA / 'office-pdfs'
    office_pdfs.mkdir(exist_ok=True)
    converted = render_office_batch(office)
    if set(converted) != set(office):
        raise RuntimeError('Mindestens eine Office-Quelle ist nicht gerendert.')
    for source, data in converted.items():
        (office_pdfs / (source.stem + '.pdf')).write_bytes(data)
    source = next(CASE.glob('*.xlsx'))
    source_bytes = source.read_bytes()
    variants = {'native-recalc': {'xl/worksheets/sheet2.xml': {'C7': 9450}}}
    for day in (26, 27):
        serial = (datetime(2026, 9, day) - datetime(1899, 12, 30)).days
        changes = {'xl/worksheets/sheet2.xml': {'E7': serial}}
        if day == 27:
            changes['xl/worksheets/sheet3.xml'] = {'E7': serial}
        variants[f'weekend-{day}'] = changes
    copies = {}
    for name, changes in variants.items():
        changed = QA / f'{name}-input.xlsx'
        recalculation_copy(source, changed, changes)
        copies[changed] = QA / f'{name}-output.pdf'
    native = render_office_batch(list(copies))
    if set(native) != set(copies):
        raise RuntimeError('Native Neuberechnung konnte nicht geprüft werden.')
    for changed, output in copies.items():
        output.write_bytes(native[changed])
    assert source.read_bytes() == source_bytes, 'Regression darf die Originalmappe nicht ändern.'
    pdfs = sorted([p for p in sources() if p.suffix == '.pdf'] + list(office_pdfs.glob('*.pdf')))
    if TOTAL.exists():
        pdfs.append(TOTAL)
    renderer = shutil.which('pdftoppm')
    if not renderer:
        raise RuntimeError('pdftoppm fehlt im PATH.')
    inventory = {}
    images = []
    for pdf in pdfs:
        folder = QA / 'render' / pdf.stem
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob('page-*.png'):
            old.unlink()
        subprocess.run([renderer, '-scale-to', '1500', '-png', str(pdf), str(folder / 'page')], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
        pages = sorted(folder.glob('page-*.png'))
        inventory[pdf.name] = len(PdfReader(pdf).pages)
        assert len(pages) == inventory[pdf.name]
        if pdf != TOTAL:
            images.extend(pages)
    (QA / 'page-counts.json').write_text(json.dumps(inventory, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    for old in QA.glob('contact-*.png'):
        old.unlink()
    for start in range(0, len(images), 6):
        contact = Image.new('RGB', (1500, 1530), 'white')
        draw = ImageDraw.Draw(contact)
        for offset, image in enumerate(images[start:start + 6]):
            panel = Image.open(image).convert('RGB')
            panel.thumbnail((490, 705))
            x, y = offset % 3 * 500, offset // 3 * 765
            contact.paste(panel, (x, y + 30))
            draw.text((x + 5, y + 4), image.parent.name[:32] + ' / ' + image.stem, font=screen_font(13), fill='black')
        contact.save(QA / f'contact-{start // 6 + 1:02}.png')
    print('Renderinventar:', json.dumps(inventory, ensure_ascii=False))


class CaseTests(unittest.TestCase):
    def test_01_inventory_and_export_boundaries(self):
        paths = sources()
        self.assertEqual([int(p.name[:2]) for p in paths], list(range(1, 27)))
        self.assertEqual(Counter(p.suffix for p in paths), {'.pdf': 11, '.docx': 4, '.eml': 5, '.txt': 2, '.csv': 2, '.xlsx': 1, '.png': 1})
        self.assertFalse((CASE / '.qa').exists())
        self.assertEqual({p.name for p in CASE.iterdir()}, {p.name for p in paths} | {'README.md', 'rubric.yaml', 'gesamt-pdf'})
        for p in CASE.rglob('*'):
            if p.is_file() and ('.qa' in p.parts or p.name in {'rubric.yaml', 'README.md'}):
                self.assertFalse(include_in_working_dump(p, CASE), p.name)

    def test_02_source_quality_and_no_solutions(self):
        for p in sources():
            if p.suffix == '.png':
                with Image.open(p) as im:
                    self.assertEqual(im.size, (1480, 1000))
                continue
            text = QUALITY.export_text(p)
            self.assertFalse(SOLUTIONS.problems(p, text), p.name)
            self.assertNotIn('Benutzung auf eigene Verantwortung', text)
            self.assertNotIn('This test case file was generated', text)
            self.assertIsNone(QUALITY.SYNTHETIC_EMAIL_PATTERN.search(text), p.name)
            for label, pattern in QUALITY.EXPORT_META_PATTERNS.items():
                self.assertIsNone(pattern.search(text), (p.name, label))
            self.assertEqual(QUALITY.language_prose_errors(text, p), [], p.name)
            self.assertNotRegex(text, r'\[[^\]]*(?:Name|Betrag|Datum|TODO)[^\]]*\]')
            if p.suffix in {'.pdf', '.docx', '.eml', '.txt'}:
                self.assertGreater(len(text), 600, p.name)
            if p.suffix in {'.pdf', '.docx'}:
                self.assertIn('Sehr geehrte', text, p.name)
                self.assertIn('Bezug:', text, p.name)
                self.assertRegex(text, r'(?:freundlichen Grüßen|Freundliche Grüße|gez\.|Kenntnisnahme|Übertragung)')

    def test_03_mail_headers_and_reply(self):
        messages = []
        for p in CASE.glob('*.eml'):
            self.assertEqual(QUALITY.eml_quality_errors(p), [])
            msg = BytesParser(policy=policy.default).parsebytes(p.read_bytes())
            messages.append(msg)
            self.assertFalse(msg.defects)
            for field in ['Date', 'From', 'To', 'Subject', 'Message-ID', 'Return-Path', 'Received', 'MIME-Version', 'Content-Type', 'Content-Transfer-Encoding']:
                self.assertTrue(msg[field], (p.name, field))
            self.assertEqual(msg.get_content_charset(), 'utf-8')
            date = parsedate_to_datetime(msg['Date'])
            self.assertEqual(date.year, 2026)
            self.assertEqual(date.month, 9)
            self.assertLessEqual(date.day, 23)
        ids = {m['Message-ID'] for m in messages}
        self.assertEqual(len(ids), 5)
        for msg in messages:
            if msg['In-Reply-To']:
                self.assertIn(msg['In-Reply-To'], ids)
                self.assertEqual(msg['References'], msg['In-Reply-To'])

    def test_04_csv_and_amounts(self):
        for p in CASE.glob('*.csv'):
            self.assertEqual(CSV.validate_file(p), [])
            with p.open(encoding='utf-8-sig', newline='') as handle:
                table = list(csv.reader(handle, delimiter=';'))
            self.assertGreaterEqual(len(table) - 1, 4)
            self.assertEqual(len({len(row) for row in table}), 1)
            self.assertTrue(all(line.startswith('"') and line.endswith('"') for line in p.read_text(encoding='utf-8-sig').splitlines()))
        debtors, creditors = rows('19_'), rows('20_')
        self.assertEqual(sum(money(r['Offen_EUR']) for r in debtors), Decimal('48355'))
        self.assertEqual(sum(money(r['Offen_EUR']) for r in creditors), Decimal('41777.80'))
        for row in debtors:
            self.assertEqual(money(row['Brutto_EUR']) - money(row['Bezahlt_EUR']), money(row['Offen_EUR']))
        self.assertEqual(Decimal('120000') * Decimal('1.19'), Decimal('142800'))
        self.assertEqual(Decimal('48000') * Decimal('1.19'), Decimal('57120'))
        self.assertEqual(6 * 1900 + 24 * 325 + 6 * 900, 24600)
        self.assertEqual(Decimal('24600') * Decimal('1.19'), Decimal('29274'))
        self.assertEqual(5000 + 24274, 29274)

    def test_05_workbook_values_formulas_and_sources(self):
        p = next(CASE.glob('*.xlsx'))
        self.assertEqual(uncached_formula_cells(p), [])
        values = load_workbook(p, data_only=True)
        formulas = load_workbook(p, data_only=False)
        try:
            self.assertEqual(values.sheetnames, ['Zahlungsstand', 'Eingänge', 'Ausgänge'])
            expected = {'B9': 25749.4, 'C23': 18240, 'D23': 98407.8, 'E21': -94418.4, 'F21': -54418.4}
            for address, number in expected.items():
                self.assertAlmostEqual(values['Zahlungsstand'][address].value, number, places=6)
            self.assertEqual(values['Eingänge']['C11'].value, 48355)
            self.assertAlmostEqual(values['Ausgänge']['C18'].value, 41777.8)
            self.assertIsNone(values['Eingänge']['E6'].value)
            self.assertEqual(formulas['Zahlungsstand']['B13'].value, '=E12')
            status = values['Zahlungsstand']
            days = [status[f'A{row}'].value for row in range(12, 22)]
            self.assertEqual(days, [datetime(2026, 9, 23) + timedelta(days=i) for i in range(10)])
            for row in (15, 16):
                self.assertEqual(status[f'C{row}'].value, 0)
                self.assertEqual(status[f'D{row}'].value, 0)
                self.assertAlmostEqual(status[f'E{row}'].value, status[f'B{row}'].value)
            self.assertEqual(status['C23'].value, values['Eingänge']['C14'].value)
            self.assertAlmostEqual(status['D23'].value, values['Ausgänge']['C19'].value)
            self.assertEqual(formulas['Zahlungsstand']['C23'].value, '=SUM(C12:C21)')
            self.assertEqual(formulas['Zahlungsstand']['D23'].value, '=SUM(D12:D21)')
            self.assertEqual(status['E24'].value, status['E21'].value)
            self.assertEqual(status['F24'].value, status['F21'].value)
            for index, record in enumerate(rows('19_'), 6):
                self.assertEqual(values['Eingänge'][f'A{index}'].value, record['Beleg'])
                self.assertAlmostEqual(values['Eingänge'][f'C{index}'].value, float(money(record['Offen_EUR'])))
            for index, record in enumerate(rows('20_'), 6):
                self.assertEqual(values['Ausgänge'][f'A{index}'].value, record['Beleg'])
                self.assertAlmostEqual(values['Ausgänge'][f'C{index}'].value, float(money(record['Offen_EUR'])))
            for ws in formulas:
                self.assertTrue(ws.print_area)
                for row in ws:
                    for cell in row:
                        if cell.data_type == 'f':
                            cached = values[ws.title][cell.coordinate]
                            self.assertIsNotNone(cached.value, (ws.title, cell.coordinate))
                            self.assertNotEqual(cached.data_type, 'e')
                            self.assertNotIn('[', cell.value)
            self.assertIsInstance(values['Zahlungsstand']['A12'].value, datetime)
            self.assertEqual(values['Zahlungsstand']['A12'].value.date().isoformat(), '2026-09-23')
        finally:
            values.close()
            formulas.close()

    def test_06_docx_typography_and_real_blank_lines(self):
        for p in list(CASE.glob('*.docx')) + list((QA / 'office').glob('*.docx')):
            d = Document(p)
            self.assertTrue(QUALITY.is_a4(d), p.name)
            self.assertEqual(d.styles['Normal'].font.name, 'Times New Roman')
            self.assertEqual(d.styles['Normal'].font.size.pt, 11)
            paragraphs = d.paragraphs
            for index, paragraph in enumerate(paragraphs):
                if paragraph.style.style_id in {'Heading1','Heading2'}:
                    self.assertRegex(paragraph.text, r'^\d+(?:\.\d+)* ')
                    self.assertEqual(paragraphs[index + 1].text, '')
                    self.assertTrue(paragraphs[index + 1].paragraph_format.keep_with_next)

    def test_07_readme_and_rubric(self):
        errors = []
        self.assertEqual(README.validate_local_readme(CASE.name, CASE, errors), 6)
        self.assertEqual(errors, [])
        text = (CASE / 'README.md').read_text()
        self.assertEqual(text.count('Diese Testakte wurde mit KI generiert'), 1)
        rubric = (CASE / 'rubric.yaml').read_text()
        self.assertEqual(rubric.count('check_type: human_review'), 11)
        self.assertNotIn('rubric.yaml', '\n'.join(p.name for p in sources()))

    def test_08_pdf_readability_and_total(self):
        pdfs = list(CASE.glob('*.pdf')) + [TOTAL]
        for p in pdfs:
            self.assertTrue(p.exists(), p.name)
            self.assertTrue(QUALITY.pdf_is_a4(p), p.name)
            reader = PdfReader(p)
            for index, page in enumerate(reader.pages):
                text = page.extract_text() or ''
                self.assertTrue(text.strip() or page.images, (p.name, index + 1))
                self.assertNotIn('Benutzung auf eigene Verantwortung', text)
                self.assertNotIn('This test case file was generated', text)
                self.assertNotIn('\ufffd', text)
        text = QUALITY.pdf_text(TOTAL)
        for p in sources():
            self.assertIn(p.name, text, p.name)
        self.assertNotIn('human_review', text)
        self.assertNotIn('rubric.yaml', text)

    def test_09_chronology_and_authority(self):
        register = QUALITY.pdf_text(CASE / '02_Register_und_Anteilsstand_2026-09-22.pdf')
        self.assertIn('Einzelvertretungsbefugnis', register)
        self.assertIn('Einzelprokura', register)
        statute = ' '.join(QUALITY.pdf_text(CASE / '03_Gesellschaftsvertrag_Auszug.pdf').split())
        self.assertIn('vierzehn volle Kalendertage', statute)
        draft = QUALITY.docx_text(CASE / '08_Einladung_Entwurf_2026-09-23.docx')
        self.assertIn('noch nicht unterschrieben oder versandt', draft)
        self.assertEqual((datetime(2026, 10, 9) - datetime(2026, 9, 23)).days - 1, 15)
        self.assertEqual(datetime(2026, 9, 25).weekday(), 4)
        self.assertEqual(datetime(2026, 10, 9).weekday(), 4)
        self.assertEqual(datetime(2026, 10, 12).weekday(), 0)
        customer = QUALITY.pdf_text(CASE / '12_Kundenbrief_Maengel_2026-09-15.pdf')
        self.assertIn('keine Abnahme', customer)
        bank = QUALITY.pdf_text(CASE / '17_Bankbrief_Kreditaufstockung_2026-09-21.pdf')
        self.assertIn('noch keine abrufbare Zusage', bank)

    def test_10_native_print_and_recalculation(self):
        printed = QA / 'office-pdfs/25_Liquiditaetsstatus_2026-09-23.pdf'
        recalculated = QA / 'native-recalc-output.pdf'
        if not printed.exists() or not recalculated.exists():
            self.skipTest('Mit --render wird die native Druck- und Neuberechnungsprüfung ausgeführt.')
        self.assertEqual(len(PdfReader(printed).pages), 3)
        text = ' '.join(QUALITY.pdf_text(printed).split())
        workbook = load_workbook(next(CASE.glob('*.xlsx')), data_only=True)
        try:
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        if isinstance(cell.value, str):
                            self.assertIn(' '.join(cell.value.split()), text, (sheet.title, cell.coordinate))
        finally:
            workbook.close()
        native = QUALITY.pdf_text(recalculated)
        self.assertIn('53,418.40', native)
        self.assertNotIn('54,418.40', native)

    def test_11_weekend_native_recalculation(self):
        outputs = [QA / f'weekend-{day}-output.pdf' for day in (26, 27)]
        if not all(p.exists() for p in outputs):
            self.skipTest('Mit --render werden Wochenendänderungen in Wegwerfkopien neu berechnet.')
        for day, output in zip((26, 27), outputs):
            with self.subTest(day=day):
                pages = PdfReader(output).pages
                self.assertEqual(len(pages), 3)
                daily = ' '.join(pages[0].extract_text().split())
                incoming = ' '.join(pages[1].extract_text().split())
                outgoing = ' '.join(pages[2].extract_text().split())
                total = re.search(r'Summe im Zeitraum\s+([\d,.]+)\s+([\d,.]+)', daily)
                scheduled = re.search(r'Terminierte Eingänge im Planzeitraum\s+([\d,.]+)', incoming)
                payments = re.search(r'Ausgänge bis 02\.10\.\s+([\d,.]+)', outgoing)
                self.assertIsNotNone(total)
                self.assertIsNotNone(scheduled)
                self.assertIsNotNone(payments)
                self.assertEqual(total.group(1), scheduled.group(1))
                self.assertEqual(total.group(1), '18,240.00')
                self.assertEqual(total.group(2), payments.group(1))
                self.assertEqual(total.group(2), '98,407.80')
                outgoing_amount = r'0\.00' if day == 26 else r'4,879\.00'
                self.assertRegex(daily, rf'2026-09-{day}\s+\([\d,.]+\)\s+8,450\.00\s+{outgoing_amount}')
                self.assertIn('54,418.40', daily)
                self.assertNotIn('62,868.40', daily)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--render', action='store_true')
    args = parser.parse_args()
    if args.render:
        render_sources()
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(CaseTests)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    report = {
        'tests_run': result.testsRun,
        'failures': len(result.failures), 'errors': len(result.errors), 'skipped': len(result.skipped),
        'sources': len(sources()), 'formats': dict(Counter(p.suffix for p in sources())),
        'combined_pdf_pages': len(PdfReader(TOTAL).pages) if TOTAL.exists() else None,
        'bank_available': 25749.4, 'debtors_open': 48355, 'creditors_open': 41777.8,
        'planned_inflows': 18240, 'planned_outflows': 98407.8,
        'planned_final_bank': -94418.4, 'planned_final_with_current_line': -54418.4,
        'release_scope': 'ZIPs und zentrale Verknüpfungen bleiben beim Elternagenten.',
    }
    (QA / 'test-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if result.wasSuccessful() else 1


if __name__ == '__main__':
    raise SystemExit(main())
