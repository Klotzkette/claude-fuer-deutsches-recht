#!/usr/bin/env python3
"""Zwei native KI-Arbeitsakten; keine Gesamt-PDFs oder Archive.

Reihenfolge: --phase base, MJS-Builder, --phase render, --phase messages,
--phase check. Der Renderlauf verwendet jeweils genau einen Office-Prozess.
"""
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email.headerregistry import Address
from email.utils import getaddresses
from email.policy import SMTP, default
from email.parser import BytesParser
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
from testakte_file_filter import include_in_working_dump
import argparse, hashlib, json, os, re, shutil, subprocess, sys, zipfile
import xml.etree.ElementTree as ET

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'scripts/data/ki-vertiefung-standards-akten.json'
QA=ROOT/'quality/ki-verordnung-2026-10-09-vertiefung'
WORK=Path(os.environ.get('STANDARDS_AKTEN_QA','/tmp/ki-standards-akten-qa'))
RENDER=Path.home()/'.codex/plugins/cache/openai-primary-runtime/documents/26.1007.11041/skills/documents/render_docx.py'
WARNING='> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n>\n> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'

def cases(): return json.loads(DATA.read_text())

def word(item,out,case):
    d=Document(); section=d.sections[0]
    section.page_width=Cm(21); section.page_height=Cm(29.7)
    section.left_margin=section.right_margin=Cm(2.25)
    section.top_margin=Cm(1.8); section.bottom_margin=Cm(1.9)
    for st in d.styles:
        if st.type==1:
            st.font.name='Times New Roman'; st.font.size=Pt(11); st.font.color.rgb=RGBColor(0,0,0)
    for name,size in [('Title',15),('Heading 1',11)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.keep_with_next=True;st.paragraph_format.space_after=Pt(5)
        st.paragraph_format.space_before=Pt(6 if name=='Heading 1' else 0)
    d.styles['Normal'].paragraph_format.space_after=Pt(5)
    d.styles['Normal'].paragraph_format.line_spacing=1.0
    d.styles['Normal'].paragraph_format.widow_control=True
    for n in list(d.styles.element.iter()):
        if n.tag==qn('w:pBdr'):n.getparent().remove(n)
        if n.tag==qn('w:rFonts'):
            for k in list(n.attrib):
                if 'theme' in k.lower():del n.attrib[k]
            n.set(qn('w:ascii'),'Times New Roman');n.set(qn('w:hAnsi'),'Times New Roman')
    head=section.header.paragraphs[0];head.text=case['ref']+'   '+item['author'].split('|')[0].strip()
    head.style=d.styles['Normal']
    for r in head.runs:r.font.size=Pt(9)
    d.add_paragraph(item['title'],'Title')
    d.add_paragraph(item['author'].replace(' | ','\n'))
    for p in re.split(r'\n\s*\n',item['body']):
        heading=bool(re.match(r'^\d+(?:\.\d+)* [^\n]+$',p))
        d.add_paragraph(p,'Heading 1' if heading else None)
    foot=section.footer.paragraphs[0];foot.alignment=2;foot.add_run(case['ref']+'   Seite ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    for run in foot.runs:run.font.size=Pt(9)
    d.core_properties.title=item['title'];d.core_properties.author=item['author'].split('|')[0].strip()
    d.core_properties.subject=case['ref'];d.core_properties.created=d.core_properties.modified=datetime(2026,10,9,8)
    d.save(out)

def readme(c,out):
    records=[]
    for i in c['docs']:
        records.extend([(i['file'],i['title']),(i['file'].replace('.docx','.pdf'),i['title']+' als Lesefassung')])
    records.extend((i['file'],i['subject']) for i in c['emails'])
    records.extend([(c['xlsx'],'Arbeitsmappe mit Eingaben, Formeln und beschriebenen Zählwerten'),(c['chatfile'],'Zeitlich zugeordneter Projektchat')])
    rows='\n'.join(f'| [{name}]({name}) | {title} |' for name,title in sorted(records))
    text=f'''# 1. {c['title']}

## 1.1. Auftrag und Aktenstand

{c['intro']}

Aktenstand ist der 9. Oktober 2026. Einstieg sind der Prüfauftrag und die E-Mail 07. Der vorhandene Arbeitsentwurf ist anhand aller Unterlagen fortzuschreiben. Die Akte enthält 20 native Dateien: sechs Word-Dokumente, ihre sechs PDF-Lesefassungen, sechs E-Mails, eine Excel-Arbeitsmappe und einen Textchat. Alle sechs E-Mails enthalten die bezeichneten Dateien tatsächlich als bytegleiche Anhänge. Die PDF-Paare und Mailanhänge sind keine zusätzlichen unabhängigen Beweise.

Alle Personen, Einrichtungen, Anschriften und Vorgänge sind erfunden. Kontaktadressen verwenden ausschließlich reservierte .example-Domains. Aussagen von Beteiligten und offene Stellen des Arbeitsentwurfs sind zu prüfendes Material, keine Musterlösung. Tatsächlicher Versand, Registereintrag und produktive Änderung sind nicht beauftragt.

<!-- reserved-example-contacts -->

## 1.2. Native Unterlagen und Lesefassungen

{WARNING}

| Datei | Gegenstand |
| --- | --- |
{rows}
'''
    (out/'README.md').write_text(text)

def render(c,out):
    logs=[]
    for item in c['docs']:
        source=out/item['file']; dest=WORK/c['slug']/source.stem;dest.mkdir(parents=True,exist_ok=True)
        for stale in dest.glob('page-*.png'):stale.unlink()
        command=[sys.executable,str(RENDER),str(source),'--output_dir',str(dest),'--emit_pdf','--dpi','115']
        result=subprocess.run(command,env=os.environ.copy(),text=True,capture_output=True,timeout=180)
        (dest/'render.log').write_text(result.stdout+'\n'+result.stderr)
        if result.returncode:raise RuntimeError(source.name+': '+result.stderr[-3000:])
        pdf=dest/(source.stem+'.pdf')
        if not pdf.exists():raise RuntimeError('Fehlende Lesefassung '+str(pdf))
        shutil.copy2(pdf,source.with_suffix('.pdf'))
        pages=len(PdfReader(pdf).pages)
        logs.append({'file':str(source.relative_to(ROOT)),'pdf_pages':pages,'rendered_pngs':len(list(dest.glob('page-*.png')))})
        print(c['slug'],source.name,pages,'Seiten',flush=True)
    return logs

def message(item,out):
    m=EmailMessage(policy=SMTP)
    for h,k in [('From','sender'),('To','to')]:m[h]=tuple(Address(display_name=n,addr_spec=a) for n,a in getaddresses([item[k]]))
    m['Subject']=item['subject'];m['Date']=datetime.fromisoformat(item['date'])
    key=hashlib.sha256((out.name+item['file']).encode()).hexdigest()[:24]
    m['Message-ID']=f'<{key}@projektpost.example>'
    m.set_content(item['body']+'\n',charset='utf-8')
    types={'.pdf':'pdf','.docx':'vnd.openxmlformats-officedocument.wordprocessingml.document','.xlsx':'vnd.openxmlformats-officedocument.spreadsheetml.sheet'}
    for name in item['attachments']:
        f=out/name;m.add_attachment(f.read_bytes(),maintype='application',subtype=types[f.suffix],filename=name)
    m.set_boundary('projektpost-'+key);(out/item['file']).write_bytes(m.as_bytes())

def check(c,out):
    originals=sorted(p for p in out.iterdir() if p.is_file() and p.suffix in ['.docx','.pdf','.xlsx','.eml','.txt'])
    assert len(originals)==20,(c['slug'],len(originals))
    assert all(include_in_working_dump(p,out) for p in originals), 'Originaldatei wird vom zentralen Aktenfilter ausgeschlossen'
    counts={ext:sum(p.suffix==ext for p in originals) for ext in ['.docx','.pdf','.eml','.xlsx','.txt']}
    assert counts=={'.docx':6,'.pdf':6,'.eml':6,'.xlsx':1,'.txt':1},counts
    checks=[]
    for item in c['docs']:
        f=out/item['file'];d=Document(f);text='\n'.join(p.text for p in d.paragraphs)
        assert len(text)>2500,(f,len(text))
        pdf=f.with_suffix('.pdf');r=PdfReader(pdf);pdftext='\n'.join(p.extract_text() or '' for p in r.pages)
        assert len(pdftext)>2500 and 'Benutzung auf eigene Verantwortung' not in pdftext,pdf
        assert len(r.pages)<=4,(pdf,len(r.pages))
        checks.append({'docx':f.name,'characters':len(text),'pdf_pages':len(r.pages),'pdf_characters':len(pdftext)})
    attachments=[]
    for item in c['emails']:
        m=BytesParser(policy=default).parsebytes((out/item['file']).read_bytes())
        assert not m.defects,m.defects
        assert len(m.get_body(preferencelist=('plain',)).get_content())>700,item['file']
        for h in ['From','To']:
            assert all(addr.endswith('.example') for _,addr in getaddresses([str(m[h])]))
        got=[]
        for a in m.iter_attachments():
            name=a.get_filename();payload=a.get_payload(decode=True)
            assert payload==(out/name).read_bytes(),(item['file'],name)
            got.append(name);attachments.append({'email':item['file'],'attachment':name,'sha256':hashlib.sha256(payload).hexdigest()})
        assert got==item['attachments'],(got,item['attachments'])
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(out/c['xlsx']) as z:
        formula_count=0;cached=0
        for n in z.namelist():
            if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',n):
                for cell in ET.fromstring(z.read(n)).findall('.//s:c',ns):
                    if cell.find('s:f',ns) is not None:
                        formula_count+=1
                        if cell.find('s:v',ns) is not None:cached+=1
                        assert cell.get('t')!='e',(n,cell.attrib)
        assert formula_count>=10 and cached==formula_count,(formula_count,cached)
    return {'case':c['slug'],'counts':counts,'documents':checks,'attachments':attachments,'formula_count':formula_count,'cached_formula_count':cached,'file_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in originals}}

def main():
    p=argparse.ArgumentParser();p.add_argument('--phase',choices=['base','render','messages','check'],required=True);args=p.parse_args()
    QA.mkdir(parents=True,exist_ok=True);WORK.mkdir(parents=True,exist_ok=True);results=[]
    for c in cases():
        out=ROOT/'testakten'/c['slug'];out.mkdir(parents=True,exist_ok=True)
        if args.phase=='base':
            for item in c['docs']:word(item,out/item['file'],c)
            (out/c['chatfile']).write_text(c['chat']);readme(c,out)
        elif args.phase=='render':results.extend(render(c,out))
        elif args.phase=='messages':
            for item in c['emails']:message(item,out)
        elif args.phase=='check':results.append(check(c,out))
    if results:(QA/f'standards-akten-{args.phase}.json').write_text(json.dumps(results,ensure_ascii=False,indent=2)+'\n')
    print(args.phase,'abgeschlossen',flush=True)

if __name__=='__main__':main()
