#!/usr/bin/env python3
"""Sichert Umfang und Belegkonsistenz der Berliner Personalakte."""

import csv
import unittest
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
CASE = ROOT / "testakten/arbeitsrecht-asset-deal-50-vertraege-berlin"


def text(name):
    return "\n".join(p.text for p in Document(CASE / name).paragraphs)


class PersonalakteTest(unittest.TestCase):
    def test_fifty_contracts_and_two_separate_directors(self):
        files = sorted(CASE.glob("*_arbeitsvertrag_*.docx"))
        self.assertEqual(len(files), 50)
        for number, file in enumerate(files, 1):
            with self.subTest(file=file.name):
                self.assertTrue(file.name.startswith(f"{number:03}_"))
                value = text(file.name)
                self.assertIn(f"P{number:03}", value)
                self.assertGreater(len(value), 8000)
        self.assertEqual(len(list(CASE.glob("*_dienstvertrag_*.docx"))), 2)

    def test_working_files_and_readme(self):
        files = [p for p in CASE.iterdir() if p.suffix in {".docx", ".xlsx", ".eml", ".csv", ".txt"}]
        self.assertEqual(len(files), 110)
        readme = (CASE / "README.md").read_text()
        for file in files:
            with self.subTest(file=file.name):
                self.assertIn(f"]({file.name})", readme)
                self.assertRegex(file.name, r"^[a-z0-9_]+\.[a-z]+$")

    def test_native_correspondence(self):
        emails = list(CASE.glob("*.eml"))
        self.assertEqual(len(emails), 15)
        for file in emails:
            with self.subTest(file=file.name):
                message = BytesParser(policy=policy.default).parsebytes(file.read_bytes())
                for header in ("From", "To", "Date", "Subject", "Message-ID", "MIME-Version"):
                    self.assertTrue(message[header])
                self.assertFalse(message.defects)
                self.assertEqual(message["Content-Language"], "en-GB" if file.name.startswith("113_") else "de-DE")
                self.assertGreater(len(message.get_body(preferencelist=("plain",)).get_content()), 500)

    def test_csv_shape_and_bank_reconciliation(self):
        for file in CASE.glob("*.csv"):
            with self.subTest(file=file.name), file.open(encoding="utf-8-sig", newline="") as source:
                rows = list(csv.reader(source, delimiter=";"))
                self.assertGreaterEqual(len(rows), 5)
                self.assertTrue(all(len(row) == len(rows[0]) for row in rows))
        with (CASE / "118_kontoumsaetze_august.csv").open(encoding="utf-8-sig", newline="") as source:
            rows = list(csv.DictReader(source, delimiter=";"))
        balance = Decimal(rows[0]["Saldo EUR"])
        for row in rows[1:]:
            balance += Decimal(row["Gutschrift EUR"]) - Decimal(row["Belastung EUR"])
            self.assertEqual(balance, Decimal(row["Saldo EUR"]))
        self.assertEqual(balance, Decimal("42900"))

    def test_personnel_and_payroll(self):
        path = CASE / "122_personal_und_abrechnung.xlsx"
        values = load_workbook(path, data_only=True)
        formulas = load_workbook(path, data_only=False)
        try:
            expected = [f"P{i:03}" for i in range(1, 51)]
            for name in ("Personalbestand", "Augustabrechnung", "Arbeitsbedingungen"):
                self.assertEqual([values[name].cell(i, 1).value for i in range(6, 56)], expected)
            sheet = values["Augustabrechnung"]
            self.assertIn("[$-407]", formulas["Augustabrechnung"]["E6"].number_format)
            self.assertIn("[$-407]", formulas["Arbeitsbedingungen"]["E6"].number_format)
            self.assertAlmostEqual(sheet["E56"].value, 152079.56, places=2)
            self.assertEqual(sheet["F56"].value, 49000)
            self.assertAlmostEqual(sheet["G56"].value, 103079.56, places=2)
            for row in range(6, 56):
                self.assertEqual(formulas["Augustabrechnung"].cell(row, 7).value, f"=E{row}-F{row}")
                self.assertAlmostEqual(sheet.cell(row, 7).value, sheet.cell(row, 5).value - sheet.cell(row, 6).value, places=2)
            self.assertEqual(values["Arbeitsbedingungen"]["E7"].value, 3400)
            self.assertEqual(values["Arbeitsbedingungen"]["E12"].value, 3300)
            self.assertEqual(sheet["E24"].value, 0)
            for worksheet in values:
                self.assertEqual(worksheet.page_setup.orientation, "landscape")
                self.assertTrue(worksheet.print_title_rows)
        finally:
            values.close()
            formulas.close()

    def test_separate_evidence_preserved(self):
        unsigned = text("022_arbeitsvertrag_jonas_winter.docx")
        self.assertIn("Unterschrift Arbeitgeber: nicht eingetragen", unsigned)
        self.assertIn("Unterschrift Arbeitnehmer: nicht eingetragen", unsigned)
        self.assertEqual(len(list(CASE.glob("09?_nachtrag_sommer_*.docx"))), 4)
        ownership = text("055_gesellschafterbestand.docx")
        self.assertIn("Stammkapital beträgt 50000 EUR", ownership)
        self.assertIn("Nennbetrag von 26000 EUR", ownership)
        self.assertIn("Nennbetrag von 24000 EUR", ownership)
        status = text("056_statusbescheid_matthias_seifert.docx")
        self.assertIn("18.10.2022", status)
        self.assertIn("Erwerbsstatus", status)
        for file in CASE.glob("*.docx"):
            with self.subTest(file=file.name):
                self.assertNotIn(chr(167), text(file.name))


if __name__ == "__main__":
    unittest.main()
