#!/usr/bin/env python3
"""Native Office-Rechnung und Rendering nur für HOAI 7 bis 9. Autor: Klotzkette."""
from __future__ import annotations

import argparse
import importlib.util
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from PIL import Image, ImageDraw
from office_process import run_office
from testakte_office_pdf import office_binary, render_office_batch
from akten_build_runtime import screen_font

HERE=Path(__file__).resolve().parent


def module(name,file):
    spec=importlib.util.spec_from_file_location(name,HERE/file); obj=importlib.util.module_from_spec(spec); sys.modules[name]=obj; spec.loader.exec_module(obj); return obj


B=module("hoai789_originals","build-bauwirtschaft-hoai-7-9-akten.py")
N="{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
DC="{http://purl.org/dc/elements/1.1/}"
CORE="{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"
SPECS={
    7:dict(number=6,sheet="Planer-LV",area="A1:F26",cells={"F5":39736,"F7":47285.84},mutation=("C11",13,"F5",40686)),
    8:dict(number=12,sheet="Kostenjournal",area="A1:G27",cells={"D5":394586,"E6":377840,"G7":448964.6},mutation=("E12",39740,"G7",449083.6)),
    9:dict(number=12,sheet="Sicherungen",area="A1:F24",cells={"D10":4641,"D11":2844.1,"F5":7485.1},mutation=("E10",1000,"F5",6485.1)),
}


def patch_xlsx(path,spec,mutation=None,clear=False):
    """Drucklayout/Metadaten; Eingabemutation ausschließlich in temporärer Testkopie."""
    out=io.BytesIO()
    with zipfile.ZipFile(path) as src,zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data=src.read(item)
            if item.filename=="docProps/core.xml":
                root=ET.fromstring(data)
                for tag in [DC+"creator",CORE+"lastModifiedBy"]:
                    element=root.find(tag)
                    if element is None:element=ET.SubElement(root,tag)
                    element.text=B.AUTHOR
                data=ET.tostring(root,encoding="utf-8",xml_declaration=True)
            elif item.filename=="xl/workbook.xml":
                root=ET.fromstring(data); names=root.find(N+"definedNames")
                if names is None:names=ET.SubElement(root,N+"definedNames")
                for el in list(names):
                    if el.get("name")=="_xlnm.Print_Area":names.remove(el)
                el=ET.SubElement(names,N+"definedName",name="_xlnm.Print_Area",localSheetId="0");el.text=f"'{spec['sheet']}'!{spec['area']}"
                calc=root.find(N+"calcPr")
                if calc is None:calc=ET.SubElement(root,N+"calcPr")
                calc.attrib.update(calcMode="auto",fullCalcOnLoad="1",forceFullCalc="1")
                data=ET.tostring(root,encoding="utf-8",xml_declaration=True)
            elif item.filename=="xl/worksheets/sheet1.xml":
                root=ET.fromstring(data); setup=root.find(N+"pageSetup")
                if setup is None:setup=ET.SubElement(root,N+"pageSetup")
                setup.attrib.update(paperSize="9",orientation="landscape",fitToWidth="1",fitToHeight="1")
                margins=root.find(N+"pageMargins")
                if margins is None:margins=ET.SubElement(root,N+"pageMargins")
                margins.attrib.update(left="0.3",right="0.3",top="0.3",bottom="0.3",header="0.1",footer="0.1")
                prop=root.find(N+"sheetPr")
                if prop is None:prop=ET.Element(N+"sheetPr");root.insert(0,prop)
                ps=prop.find(N+"pageSetUpPr")
                if ps is None:ps=ET.SubElement(prop,N+"pageSetUpPr")
                ps.set("fitToPage","1")
                if mutation:
                    cell=root.find(f".//{N}c[@r='{mutation[0]}']")
                    if cell is None:raise ValueError(mutation)
                    for child in list(cell):cell.remove(child)
                    cell.attrib.pop("t",None);ET.SubElement(cell,N+"v").text=str(mutation[1])
                if clear or mutation:
                    for cell in root.findall(f".//{N}c"):
                        if cell.find(N+"f") is not None:
                            for cache in cell.findall(N+"v"):cell.remove(cache)
                data=ET.tostring(root,encoding="utf-8",xml_declaration=True)
            dst.writestr(item,data)
    path.write_bytes(out.getvalue())


def formulas(path):
    wb=load_workbook(path,data_only=False)
    result={(s.title,c.coordinate):c.value for s in wb for row in s for c in row if c.data_type=="f"};wb.close();return result


def roundtrip(paths,folder):
    output=folder/"output";output.mkdir(parents=True,exist_ok=True)
    binary=office_binary()
    if not binary:raise RuntimeError("SOFFICE oder portable Office-Laufzeit fehlt")
    cmd=[binary,f"-env:UserInstallation={(folder/'profile').resolve().as_uri()}","--headless","--convert-to","xlsx","--outdir",str(output),*[str(p) for p in paths]]
    rc,logs=run_office(cmd,env=os.environ.copy(),timeout=180);(folder/"office.log").write_text(logs,encoding="utf-8")
    files=[output/p.name for p in paths]
    if rc or any(not f.exists() for f in files):raise RuntimeError(logs)
    return files


def recalculate(qa,phases):
    paths=[]; original={}
    for phase in phases:
        path=B.doc_path(B.FILES[phase][SPECS[phase]["number"]]);patch_xlsx(path,SPECS[phase],clear=True)
        paths.append(path);original[phase]=formulas(path)
    output=roundtrip(paths,qa/"native-baseline")
    records=[];mutations=[]
    for phase,path,source in zip(phases,output,paths):
        spec=SPECS[phase];assert formulas(path)==original[phase],f"Formeländerung Phase {phase}"
        wb=load_workbook(path,data_only=True);sheet=wb[spec["sheet"]]
        for cell,value in spec["cells"].items():assert abs(sheet[cell].value-value)<.001,(phase,cell,sheet[cell].value,value)
        assert all(c.data_type!="e" for s in wb for row in s for c in row);wb.close()
        shutil.copyfile(path,source);patch_xlsx(source,spec)
        mut=qa/"mutation-input"/path.name;mut.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(source,mut);patch_xlsx(mut,spec,mutation=spec["mutation"])
        mutations.append(mut);records.append(dict(phase=phase,file=source.name,formulas=len(original[phase]),baseline=spec["cells"]))
    for phase,path,record in zip(phases,roundtrip(mutations,qa/"native-mutation"),records):
        spec=SPECS[phase];wb=load_workbook(path,data_only=True);cell=spec["mutation"][2];value=wb[spec["sheet"]][cell].value;wb.close()
        assert abs(value-spec["mutation"][3])<.001,(phase,value)
        assert formulas(path)==original[phase]
        record["mutation"]=dict(input_cell=spec["mutation"][0],input_value=spec["mutation"][1],result_cell=cell,result=value)
    (qa/"native-recalculation.json").write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding="utf-8")
    print("Native Basis- und Mutationsläufe erfolgreich",flush=True)


def normalize(data,title):
    reader=PdfReader(io.BytesIO(data));writer=PdfWriter()
    for page in reader.pages:writer.add_page(page)
    writer.add_metadata({"/Author":B.AUTHOR,"/Creator":B.AUTHOR,"/Producer":B.AUTHOR,"/Title":title})
    out=io.BytesIO();writer.write(out);return out.getvalue()


def render(qa,phases):
    E=module("hoai789_singlepdf","build-testakten-einzelpdf-zips.py")
    renderer=os.environ.get("DOCX_RENDERER")
    if not renderer or not Path(renderer).is_file():raise RuntimeError("DOCX_RENDERER muss auf render_docx.py der Dokumentwerkzeuge zeigen")
    poppler=shutil.which("pdftoppm")
    if not poppler:raise RuntimeError("pdftoppm aus Poppler im PATH benötigt")
    paths=[B.doc_path(d) for phase in phases for d in B.FILES[phase].values() if d["filename"].endswith(".xlsx")]
    office=render_office_batch(paths);index=[]
    for phase in phases:
        for d in B.FILES[phase].values():
            source=B.doc_path(d);ext=source.suffix.lower();base=qa/f"phase-{phase}"/f"{d['number']:02d}";base.mkdir(parents=True,exist_ok=True)
            if ext==".docx":
                command=[os.environ.get("DOCX_PYTHON",sys.executable),renderer,str(source),"--output_dir",str(base/"docx"),"--emit_pdf","--dpi","100"]
                result=subprocess.run(command,capture_output=True,text=True)
                if result.returncode:raise RuntimeError(result.stdout+result.stderr)
                data=(base/"docx"/(source.stem+".pdf")).read_bytes()
            else:data=E.render_document_pdf(source,source.parent,office)
            if not data:raise RuntimeError(f"PDF fehlt: {source.name}")
            dest=base/(source.stem+".pdf");dest.write_bytes(normalize(data,d["title"]))
            subprocess.run([poppler,"-png","-r","100",str(dest),str(base/"page")],check=True,capture_output=True,text=True)
            pages=list(PdfReader(dest).pages);images=sorted(base.glob("page-*.png"))
            if ext==".xlsx":
                assert len(pages)==1,(source.name,"Tabelle auf mehrere Druckseiten verteilt")
                expected={7:"39.736,00",8:"448.964,60",9:"7.485,10"}[phase]
                assert expected in (pages[0].extract_text() or ""),(source.name,"Deutsche Geldnotation fehlt")
            images=[base/f"page-{i:0{len(str(len(pages)))}d}.png" for i in range(1,len(pages)+1)]
            assert all(p.is_file() for p in images)
            record=dict(phase=phase,number=d["number"],file=source.name,source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),pdf=str(dest),pages=len(pages),images=[str(p) for p in images]);index.append(record)
            print(f"LPH {phase} / {source.name}: {len(pages)} Seiten",flush=True)
    (qa/"render-index.json").write_text(json.dumps(index,ensure_ascii=False,indent=2),encoding="utf-8")
    # Kontaktbögen dienen der Navigation; lesbare Einzelbilder bleiben vollständig erhalten.
    for phase in phases:
        pics=[(d["file"],p) for d in index if d["phase"]==phase for p in d["images"]]
        for offset in range(0,len(pics),12):
            subset=pics[offset:offset+12];canvas=Image.new("RGB",(1600,4*580),"#d7d7d7");draw=ImageDraw.Draw(canvas)
            for n,(name,image) in enumerate(subset):
                im=Image.open(image).convert("RGB");im.thumbnail((380,530));x=(n%4)*400+(400-im.width)//2;y=(n//4)*580+35
                canvas.paste(im,(x,y));draw.text(((n%4)*400+9,(n//4)*580+7),name[:36],font=screen_font(14),fill="black")
            canvas.save(qa/f"contact-{phase}-{offset//12+1}.jpg",quality=90)
    return index


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("--qa",type=Path);parser.add_argument("--phase",type=int,choices=[7,8,9]);parser.add_argument("--stage",choices=["all","recalculate","render"],default="all")
    args=parser.parse_args();qa=args.qa or Path(tempfile.mkdtemp(prefix="hoai-7-9-qa-"));qa.mkdir(parents=True,exist_ok=True)
    B.define_all();B.register_fonts();phases=[args.phase] if args.phase else [7,8,9]
    if args.stage in ["all","recalculate"]:recalculate(qa,phases)
    if args.stage in ["all","render"]:render(qa,phases)
    print(f"QA-Ausgabe: {qa}")


if __name__=="__main__":main()
