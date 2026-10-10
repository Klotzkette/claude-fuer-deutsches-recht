#!/usr/bin/env python3
"""Wirkungsprüfungen des lokalen VVT-Helfers ohne Netzwerk oder Echtdaten."""
from __future__ import annotations
import copy
from datetime import date, datetime
import importlib.util
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'verarbeitungsverzeichnis/scripts/vvt.py'
spec = importlib.util.spec_from_file_location('vvt', SCRIPT)
vvt = importlib.util.module_from_spec(spec)
spec.loader.exec_module(vvt)

class RuntimeTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='vvt-test-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / 'register.json'
        self.actor = 'Alma Gundelfinger'
        self.org = {'name':'Federkiel GmbH', 'contact':'Klara Gundelfinger, info@example.org', 'representative':'Nicht einschlägig nach dokumentierter Prüfung.', 'dpo':'Bestellung wird geprüft.'}
        self.doc = vvt.create(self.path, self.org, self.actor, 'Übernahme des Verzeichnisses')
        self.patch = {'id':'VT-0001', 'title':'Lohnabrechnung', 'owner':'Alma Gundelfinger', 'purpose':'Vergütung berechnen und gesetzliche Nachweise führen.',
                      'data_subjects':'Beschäftigte', 'data_categories':'Stamm- und Vergütungsdaten', 'recipients':'Steuerkanzlei und Sozialversicherungsträger',
                      'transfers':'Keine Übermittlung außerhalb EWR nach dokumentierter Systemprüfung.', 'retention':'Löschfristen sind nach Dokumenttyp zu klären.', 'toms':'TOM-2026, Zugriffsbeschränkung und Mehrfaktorzugang.',
                      'screening':{'high_risk':'nein','art35_3':'nein','positive_list':'nein','rationale':'Einzelfallprüfung der Lohnabrechnung.','sources':'Datenflussplan Version 3'}}
        self.change([self.patch])

    def current(self):
        return vvt.read_register(self.path)

    def change(self, patches):
        d = self.current()
        return vvt.upsert(self.path, patches, d['revision'], vvt.digest(d), self.actor, 'Dokumentierter Sachstandsabgleich')

    def make_xlsx(self):
        p = self.root / 'bearbeitung.xlsx'
        vvt.export_xlsx(self.current(), p)
        return p

    def edit_xlsx(self, p, callback):
        from openpyxl import load_workbook
        wb = load_workbook(p)
        callback(wb)
        wb.save(p)

    def import_xlsx(self, p):
        return vvt.import_xlsx(self.path, p, self.actor, 'Rücklauf aus Fachabteilung')

    def assert_unchanged_error(self, action):
        before = self.path.read_bytes()
        history = sorted((self.root / 'register.json.history').glob('*'))
        with self.assertRaises(vvt.VVTError):
            action()
        self.assertEqual(before, self.path.read_bytes())
        self.assertEqual(history, sorted((self.root / 'register.json.history').glob('*')))
        self.assertFalse(Path(str(self.path)+'.lock').exists())

    def test_xlsx_roundtrip_preserves_blank_unknown_and_unicode(self):
        self.change([{'id':'VT-0001','joint_controller':'','processor_contracts':'unbekannt','systems':'Büro & Löhne <intern>'}])
        before = self.current()
        self.import_xlsx(self.make_xlsx())
        self.assertEqual(before, self.current())
        a = self.current()['activities'][0]
        self.assertEqual(a['joint_controller'], '')
        self.assertEqual(a['processor_contracts'], 'unbekannt')
        self.assertEqual(a['systems'], 'Büro & Löhne <intern>')

    def test_stale_excel_and_sha_conflict_leave_file_untouched(self):
        p = self.make_xlsx()
        self.change([{'id':'VT-0001','systems':'Neue Lohnanwendung'}])
        self.assert_unchanged_error(lambda:self.import_xlsx(p))
        d = self.current()
        self.assert_unchanged_error(lambda:vvt.upsert(self.path,[{'id':'VT-0001','owner':'Berta Vogel'}],d['revision'],'0'*64,self.actor,'Falsche Ausgangsfassung'))

    def test_duplicate_excel_ids_rejected_atomically(self):
        p = self.make_xlsx()
        self.edit_xlsx(p, lambda wb:wb['Art30'].append([c.value for c in wb['Art30'][3]]))
        self.assert_unchanged_error(lambda:self.import_xlsx(p))

    def test_failed_second_patch_does_not_apply_first_patch(self):
        self.assert_unchanged_error(lambda:self.change([{'id':'VT-0001','owner':'Ida Wolfram'},{'id':'VT-0002','title':'','role':'processor'}]))

    def test_deleting_excel_rows_never_deletes_activity(self):
        p = self.make_xlsx()
        def remove(wb):
            for name in vvt.SHEETS:
                wb[name].delete_rows(3)
        self.edit_xlsx(p, remove)
        before = self.current()
        self.import_xlsx(p)
        self.assertEqual(before, self.current())

    def test_partial_sheet_return_preserves_other_fields(self):
        p = self.make_xlsx()
        def edit(wb):
            wb['Taetigkeiten'].cell(3,5).value = 'Berta Meier'
            wb['Art30'].delete_rows(3)
            wb['Screening'].delete_rows(3)
        self.edit_xlsx(p, edit)
        old = self.current()['activities'][0]
        self.import_xlsx(p)
        new = self.current()['activities'][0]
        self.assertEqual(new['owner'], 'Berta Meier')
        self.assertEqual(new['recipients'], old['recipients'])
        self.assertEqual(new['screening'], old['screening'])

    def test_formula_strings_export_as_text_and_formulas_rejected(self):
        dangerous = '=HYPERLINK("https://example.invalid","senden")'
        self.change([{'id':'VT-0001','systems':dangerous,'legal_basis':'+SUM(1,2)','processor_contracts':'@Test'}])
        p = self.make_xlsx()
        from openpyxl import load_workbook
        wb = load_workbook(p, data_only=False)
        col = list(vvt.SHEETS['Taetigkeiten']).index('systems')+1
        self.assertEqual(wb['Taetigkeiten'].cell(3,col).data_type,'s')
        self.assertEqual(wb['Taetigkeiten'].cell(3,col).value,dangerous)
        self.import_xlsx(p)
        self.edit_xlsx(p, lambda book:setattr(book['Taetigkeiten'].cell(3,col),'value','=1+1'))
        self.assert_unchanged_error(lambda:self.import_xlsx(p))

    def test_prior_low_score_does_not_override_positive_trigger(self):
        self.change([{'id':'VT-0001','screening':{'positive_list':'ja','priority':'1'}}])
        d = self.current()
        self.assert_unchanged_error(lambda:vvt.review(self.path,'VT-0001','begruendet_nicht_erforderlich','Niedrige interne Priorität.','Interne Liste',d['revision'],vvt.digest(d),self.actor,'Prüfung'))
        vvt.review(self.path,'VT-0001','erforderlich','DSK-Listentatbestand nach geprüftem Sachverhalt erfüllt.','DSK-Liste, geprüfter Anwendungsfall',d['revision'],vvt.digest(d),self.actor,'Prüfung')
        self.assertEqual(vvt.review_state(self.current()['activities'][0]),'aktuell')

    def test_open_trigger_blocks_negative_dsfa(self):
        self.change([{'id':'VT-0001','screening':{'art35_3':'offen'}}])
        d = self.current()
        self.assert_unchanged_error(lambda:vvt.review(self.path,'VT-0001','begruendet_nicht_erforderlich','Begründung.','Normtext',d['revision'],vvt.digest(d),self.actor,'Prüfung'))

    def test_named_review_becomes_stale_on_fact_change(self):
        d = self.current()
        vvt.review(self.path,'VT-0001','begruendet_nicht_erforderlich','Begrenzte Datenverarbeitung nach Einzelprüfung ohne hohes Risiko.','Artikel 35 DSGVO und Datenflussplan',d['revision'],vvt.digest(d),self.actor,'DSFA-Screening')
        before = self.current()['activities'][0]
        self.assertEqual(vvt.review_state(before),'aktuell')
        self.change([{'id':'VT-0001','systems':'Neue KI-gestützte Mitarbeiterbewertung'}])
        after = self.current()['activities'][0]
        self.assertEqual(vvt.review_state(after),'veraltet')
        self.assertEqual(after['review']['reviewer'],self.actor)
        self.assertEqual(after['activity_revision'],before['activity_revision']+1)

    def test_import_cannot_change_review_or_organization(self):
        p = self.make_xlsx()
        self.edit_xlsx(p, lambda wb:setattr(wb['Entscheidungen'].cell(2,3),'value','erforderlich'))
        self.assert_unchanged_error(lambda:self.import_xlsx(p))
        p = self.make_xlsx()
        self.edit_xlsx(p, lambda wb:setattr(wb['Organisation'].cell(2,2),'value','Andere GmbH'))
        self.assert_unchanged_error(lambda:self.import_xlsx(p))

    def test_organization_change_reopens_reviews_and_is_logged(self):
        d = self.current()
        vvt.review(self.path,'VT-0001','erforderlich','Hohes Risiko dokumentiert.','Fachprüfung',d['revision'],vvt.digest(d),self.actor,'Prüfung')
        d = self.current()
        vvt.update_organization(self.path,{'dpo':'Dr. Martha Wunderlich, datenschutz@example.org'},d['revision'],vvt.digest(d),self.actor,'Datenschutzfunktion wechselt')
        changed = self.current()
        self.assertEqual(vvt.review_state(changed['activities'][0]),'offen')
        self.assertEqual(changed['history'][-1]['operation'],'update-org')
        self.assertEqual(changed['history'][-1]['actor'],self.actor)

    def test_processor_fields_and_controller_fields_differ(self):
        self.change([{'id':'VT-0002','role':'processor','title':'Support im Kundenauftrag','owner':'Alma Gundelfinger','controller_contacts':'Auftraggeber A, Musterweg 7, Leipzig.','processing_categories':'Für Auftraggeber A: Hosting und Backup.','transfers':'Kein Drittland nach Prüfung.','toms':'AV-TOM Anlage 3.'}])
        result = vvt.findings(self.current())[1]
        self.assertEqual(result['gaps'],[])
        self.assertNotIn('Zwecke',result['gaps'])

    def test_immutable_history_preserves_previous_exact_canonical_bytes(self):
        before = self.current()
        data = self.path.read_bytes()
        self.change([{'id':'VT-0001','status':'beendet'}])
        archive = self.root/'register.json.history'/f"r{before['revision']:06d}-{vvt.digest(before)}.json"
        self.assertEqual(archive.read_bytes(),data)
        log = self.current()['history'][-1]
        self.assertEqual(log['previous_sha256'],vvt.digest(before))
        self.assertEqual(log['changed_ids'],['VT-0001'])

    def test_missing_actor_reason_and_machine_name_rejected(self):
        d = self.current()
        for actor,reason in [('Alma','Anlass'),('Codex Agent','Anlass'),(self.actor,'')]:
            self.assert_unchanged_error(lambda:vvt.upsert(self.path,[{'id':'VT-0001','owner':'Berta Meier'}],d['revision'],vvt.digest(d),actor,reason))

    def test_xml_roundtrip_and_safe_unicode_edit(self):
        p = self.root/'register.xml'
        before = self.current()
        vvt.export_xml(before,p)
        vvt.import_xml(self.path,p,self.actor,'Unveränderter XML-Rücklauf')
        self.assertEqual(before,self.current())
        root = ET.parse(p).getroot()
        doc = json.loads(root[0].text)
        doc['activities'][0]['systems'] = 'Änderung & <lokal> ohne Netz'
        root[0].text = json.dumps(doc,ensure_ascii=False)
        ET.ElementTree(root).write(p,encoding='utf-8',xml_declaration=True)
        vvt.import_xml(self.path,p,self.actor,'XML-Sachänderung')
        self.assertEqual(self.current()['activities'][0]['systems'],'Änderung & <lokal> ohne Netz')

    def test_xml_entity_and_utf16_bypass_rejected(self):
        p = self.root/'unsafe.xml'
        examples = [b'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><x>&xxe;</x>',
                    '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x><x/>'.encode('utf-16'),
                    b'<?xml version="1.0" encoding="iso-8859-1"?><x/>']
        for raw in examples:
            p.write_bytes(raw)
            self.assert_unchanged_error(lambda:vvt.import_xml(self.path,p,self.actor,'Ungeprüfter Rücklauf'))

    def test_new_excel_record_requires_main_sheet_and_revision_zero(self):
        p = self.make_xlsx()
        def add(wb):
            a = vvt.empty_activity('VT-0002')
            a['title']='Lieferantenkontakte'
            wb['Taetigkeiten'].append([a[k] for k in vvt.SHEETS['Taetigkeiten']])
        self.edit_xlsx(p,add)
        self.import_xlsx(p)
        new = self.current()['activities'][1]
        self.assertEqual(new['activity_revision'],1)
        self.assertEqual(new['screening']['high_risk'],'offen')

    def test_active_register_lock_prevents_write(self):
        lock = Path(str(self.path)+'.lock')
        lock.write_text('anderer Vorgang')
        before = self.path.read_bytes()
        with self.assertRaises(vvt.VVTError):
            self.change([{'id':'VT-0001','owner':'Berta Meier'}])
        self.assertEqual(before,self.path.read_bytes())
        self.assertEqual(lock.read_text(),'anderer Vorgang')

    def test_duplicate_json_keys_are_rejected(self):
        p = self.root/'duplicate.json'
        p.write_text('{"id":"VT-0001","id":"VT-0002"}')
        with self.assertRaises(vvt.VVTError):
            vvt.read_json(p)

    def test_docx_exports_role_specific_full_sentences_tnr11(self):
        from docx import Document
        p = self.root/'register.docx'
        vvt.export_docx(self.current(),p)
        doc = Document(p)
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('VT-0001',text)
        self.assertIn('Dokumentiert ist folgende Angabe:',text)
        self.assertIn('noch nicht erhoben',text)
        self.assertEqual(doc.styles['Normal'].font.name,'Times New Roman')
        self.assertEqual(doc.styles['Normal'].font.size.pt,11)
        self.assertFalse(any(re.match(r'^[IVX]+\\.',p.text) for p in doc.paragraphs))

    def test_html_escapes_markup_and_embedded_script_close(self):
        attack='</script><script>alert("hi")</script>'
        self.change([{'id':'VT-0001','title':attack}])
        p = self.root/'register.html'
        vvt.export_html(self.current(),p)
        text=p.read_text()
        self.assertNotIn(attack,text)
        self.assertIn('\\u003c/script',text)
        payload=re.search(r'<script type="application/json" id="data">(.*?)</script>',text,re.S).group(1)
        self.assertEqual(json.loads(payload)['activities'][0]['title'],attack)
        self.assertNotRegex(text,r'<script[^>]+src=')

    def test_native_excel_dates_roundtrip_and_edit(self):
        from openpyxl import load_workbook
        self.change([{'id':'VT-0001','next_review':'2027-03-15'}])
        p=self.make_xlsx()
        col=list(vvt.SHEETS['Taetigkeiten']).index('next_review')+1
        wb=load_workbook(p)
        self.assertIsInstance(wb['Taetigkeiten'].cell(3,col).value,(date,datetime))
        before=self.current()
        self.import_xlsx(p)
        self.assertEqual(before,self.current())
        def edit(book):
            c=book['Taetigkeiten'].cell(3,col)
            c.value=datetime(2027,4,20)
            c.number_format='dd.mm.yyyy'
        self.edit_xlsx(p,edit)
        self.import_xlsx(p)
        self.assertEqual(self.current()['activities'][0]['next_review'],'2027-04-20')

    def test_excel_date_numbers_and_times_are_not_guessed(self):
        col=list(vvt.SHEETS['Taetigkeiten']).index('next_review')+1
        for value,fmt in [(45000,'General'),(datetime(2027,3,15,12,30),'dd.mm.yyyy hh:mm'),('2027-02-30','@')]:
            p=self.make_xlsx()
            def edit(book):
                c=book['Taetigkeiten'].cell(3,col)
                c.value=value
                c.number_format=fmt
            self.edit_xlsx(p,edit)
            self.assert_unchanged_error(lambda:self.import_xlsx(p))

    def test_xml_adds_new_activity_and_control_characters_are_blocked(self):
        p=self.root/'register.xml'
        vvt.export_xml(self.current(),p)
        root=ET.parse(p).getroot()
        doc=json.loads(root[0].text)
        a=vvt.empty_activity('VT-0002')
        a.update(title='Besucherempfang',activity_revision=1)
        doc['activities'].append(a)
        root[0].text=json.dumps(doc,ensure_ascii=False)
        ET.ElementTree(root).write(p,encoding='utf-8',xml_declaration=True)
        vvt.import_xml(self.path,p,self.actor,'Tätigkeit ergänzt')
        self.assertEqual(len(self.current()['activities']),2)
        self.assert_unchanged_error(lambda:self.change([{'id':'VT-0001','systems':'A'+chr(0)+'B'}]))

    def test_planned_activity_review_does_not_activate_it(self):
        self.change([{'id':'VT-0001','status':'geplant'}])
        d=self.current()
        vvt.review(self.path,'VT-0001','erforderlich','Neue geplante KI-Prüfung erfordert DSFA.','Geprüfter Projektplan',d['revision'],vvt.digest(d),self.actor,'Prüfung vor Einführung')
        self.assertEqual(self.current()['activities'][0]['status'],'geplant')
        p=self.make_xlsx()
        self.import_xlsx(p)
        self.assertEqual(self.current()['activities'][0]['status'],'geplant')
        from openpyxl import load_workbook
        wb=load_workbook(p)
        self.assertTrue(any('geplant' in dv.formula1 for dv in wb['Taetigkeiten'].data_validations.dataValidation))

    def test_cli_status_and_invalid_write_exit(self):
        result=subprocess.run([sys.executable,str(SCRIPT),'status',str(self.path)],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        data=json.loads(result.stdout)
        self.assertEqual(data['activities'],1)
        self.assertEqual(data['sha256'],vvt.digest(self.current()))
        result=subprocess.run([sys.executable,str(SCRIPT),'validate',str(self.root/'missing.json')],capture_output=True,text=True)
        self.assertEqual(result.returncode,2)
        self.assertIn('VVT:',result.stderr)

if __name__=='__main__':
    unittest.main(verbosity=2)
