#!/usr/bin/env python3
"""Erzeugt ausschließlich fünf zugewiesene Quellakten; keine zentralen Pakete."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import mimetypes
import os
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from datetime import datetime, timedelta
from decimal import Decimal
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from email.utils import format_datetime
from pathlib import Path
from xml.etree import ElementTree as ET

sys.dont_write_bytecode = True
from bau_rundum_vertiefung_daten import CASES

ROOT = Path(__file__).resolve().parents[1]
RUNTIME = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies'
NODE = RUNTIME / 'node/bin/node'
SOFFICE = RUNTIME / 'bin/override/soffice'
NOTICE = ('Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n\n'
          'This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.')
FONTDIR = Path('/System/Library/Fonts/Supplemental')
QA_ROOT = Path('/tmp/bau-rundum-vertiefung-qa')

SHEET_JS = r'''
import fs from 'node:fs/promises';
import {Workbook, SpreadsheetFile} from '@oai/artifact-tool';
let raw=''; for await (const chunk of process.stdin) raw+=chunk;
const spec=JSON.parse(raw), wb=Workbook.create();
const excelDate=s=>Math.round((Date.parse(s+'T00:00:00Z')-Date.UTC(1899,11,30))/86400000);
const verify=(sh,cell,expected)=>{const got=sh.getRange(cell).values[0][0];
 if(typeof got!=='number'||Math.abs(got-expected)>0.011) throw Error(`${sh.name}!${cell}: ${got} != ${expected}`);};
for(const s of spec.sheets) {
 const sh=wb.worksheets.add(s.name); sh.showGridLines=false;
 sh.getRange('A2:F2').merge(); sh.getRange('A2').values=[[spec.title]];
 sh.getRange('A3:F3').merge(); sh.getRange('A3').values=[[spec.owner+' | '+spec.date]];
 sh.getRange('A5:F5').values=[s.headers];
 const all=[...s.rows,...s.tail];
 for(let i=0;i<all.length;i++) for(let j=0;j<6;j++) {
   const v=all[i][j];
   if(typeof v==='string'&&!v.startsWith('=')) sh.getRange(String.fromCharCode(65+j)+(6+i)).setNumberFormat('@');
 }
 const matrix=all.map(r=>Array.from({length:6},(_,j)=>{const v=r[j]??null;return typeof v==='string'&&v.startsWith('=')?null:(v&&v.date?excelDate(v.date):v);}));
 sh.getRange(`A6:F${5+all.length}`).values=matrix;
 for(let i=0;i<all.length;i++) for(let j=0;j<6;j++) {
   const v=all[i][j], cell=String.fromCharCode(65+j)+(6+i);
   if(typeof v==='string'&&v.startsWith('=')) sh.getRange(cell).formulas=[[v]];
   if(v&&v.date) sh.getRange(cell).setNumberFormat('dd.mm.yyyy');
 }
 let end=5+all.length;
 s.notes.forEach((n,i)=>{const r=end+2+i;sh.getRange(`A${r}:F${r}`).merge();sh.getRange(`A${r}`).values=[[n]];});
 end+=1+s.notes.length;
 sh.getRange(`A2:F${end}`).format.font={name:'Arial',size:11,color:'#202020'};
 sh.getRange(`A2:F${end}`).format.rowHeight=22;
 sh.getRange(`A2:F${end}`).format.verticalAlignment='center';
 sh.getRange(`A2:F${end}`).format.wrapText=true;
 [25,32,13,15,17,22].forEach((w,i)=>sh.getRange(`${String.fromCharCode(65+i)}1:${String.fromCharCode(65+i)}${end}`).format.columnWidth=w);
 sh.getRange('A2:F2').format.font={name:'Arial',size:14,bold:true,color:'#202020'};
 sh.getRange('A2:F3').format.rowHeight=28;
 sh.getRange('A5:F5').format={fill:'#3F4D46',font:{name:'Arial',size:11,bold:true,color:'#FFFFFF'},rowHeight:38,wrapText:true};
 for(let i=6;i<=5+all.length;i++) {
  if(i%2===0) sh.getRange(`A${i}:F${i}`).format.fill='#F1F3F2';
  sh.getRange(`D${i}:F${i}`).setNumberFormat('#,##0.00');
 }
 for(let i=0;i<all.length;i++) for(let j=0;j<6;j++) if(all[i][j]&&all[i][j].date)
  sh.getRange(String.fromCharCode(65+j)+(6+i)).setNumberFormat('dd.mm.yyyy');
 for(let i=0;i<all.length;i++) if(typeof all[i][1]==='number')
  sh.getRange('B'+(6+i)).setNumberFormat(all[i][1]>0&&all[i][1]<1?'0.0%':'#,##0.00');
 sh.freezePanes.freezeRows(5);
}
wb.recalculate();
for(const s of spec.sheets) {
 const sh=wb.worksheets.getItem(s.name);
 for(const [cell,n] of Object.entries(s.controls)) verify(sh,cell,n);
 const inspection=await wb.inspect({kind:'table',range:`'${s.name}'!A5:F${5+s.rows.length+s.tail.length}`,include:'values,formulas',tableMaxRows:30,tableMaxCols:6,maxChars:12000});
 await fs.writeFile(spec.qa+'/'+s.name+'-inspection.json',inspection.ndjson);
 const preview=await wb.render({sheetName:s.name,range:`A1:F${7+s.rows.length+s.tail.length+s.notes.length}`,scale:1.5,format:'png'});
 await fs.writeFile(spec.qa+'/'+s.name+'.png',new Uint8Array(await preview.arrayBuffer()));
}
const errors=await wb.inspect({kind:'match',searchTerm:'#REF!|#DIV/0!|#VALUE!|#NAME\\?|#N/A|#NUM!|#NULL!',options:{useRegex:true,maxResults:50},maxChars:4000});
await fs.writeFile(spec.qa+'/errors.json',errors.ndjson);
const out=await SpreadsheetFile.exportXlsx(wb);await out.save(spec.path);
for(const s of spec.sheets) {
 const m=s.mutation, sh=wb.worksheets.getItem(s.name), before=sh.getRange(m.cell).values;
 sh.getRange(m.cell).values=[[m.date?excelDate(m.date):m.value]];wb.recalculate();verify(sh,m.result,m.expected);
 const changed=await SpreadsheetFile.exportXlsx(wb);await changed.save(spec.qa+'/mutation.xlsx');
 sh.getRange(m.cell).values=before;wb.recalculate();
 for(const [cell,n] of Object.entries(s.controls)) verify(sh,cell,n);
}
console.log(JSON.stringify({file:spec.path,controls:'ok',mutation:'ok'}));
'''


def dd(value):
    return datetime.fromisoformat(value[:10]).strftime('%d.%m.%Y')


def fmt(value):
    if isinstance(value, float):
        return f'{value:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return str(value)


def make_docx(path, p, project):
    project=p.get('project') or project
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.7)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    for style_name in ['Normal','Title','Heading 1','Heading 2']:
        st = d.styles[style_name]
        st.font.name, st.font.size, st.font.color.rgb = 'Times New Roman', Pt(11), RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().rFonts
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():del fonts.attrib[attr]
        for script in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+script),'Times New Roman')
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.05
    d.styles['Title'].font.size = Pt(15)
    d.styles['Heading 1'].font.bold = True
    d.styles['Heading 1'].paragraph_format.space_before = Pt(9)
    d.styles['Heading 1'].paragraph_format.space_after = Pt(8)
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    d.core_properties.author = 'Klotzkette'
    d.core_properties.title = p['title']
    d.add_paragraph(p['sender'])
    d.add_paragraph('An: '+p['recipient'])
    d.add_paragraph(f"{dd(p['date'])} | Zeichen: {p['ref']}")
    d.add_paragraph(p['title'], 'Title')
    d.add_paragraph(project)
    for heading, body in p['sections']:
        d.add_paragraph(heading, 'Heading 1')
        d.add_paragraph(body)
    if p.get('table'):
        headers, rows = p['table']
        t = d.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        for cell, value in zip(t.rows[0].cells, headers):
            cell.text = value
            shd = OxmlElement('w:shd'); shd.set(qn('w:fill'),'E7ECE9'); cell._tc.get_or_add_tcPr().append(shd)
        repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        widths = ([2,8,2,2,3] if len(headers)==5 else [2,7,3,5] if len(headers)==4 else [5,5,7])
        for row in rows:
            for cell,value in zip(t.add_row().cells,row): cell.text = fmt(value)
        for row in t.rows:
            for j,cell in enumerate(row.cells):
                cell.width = Cm(widths[j])
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    for run in paragraph.runs: run.font.size = Pt(10)
    d.add_paragraph(p['signer'])
    footer = section.footer.paragraphs[0]
    footer.add_run(p['ref']+' | Seite ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    for run in footer.runs: run.font.size = Pt(9)
    d.save(path)


def make_pdf(path, p, project):
    project=p.get('project') or project
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from xml.sax.saxutils import escape
    pdfmetrics.registerFont(TTFont('TNR', str(FONTDIR/'Times New Roman.ttf')))
    pdfmetrics.registerFont(TTFont('TNRB', str(FONTDIR/'Times New Roman Bold.ttf')))
    normal = ParagraphStyle('body',fontName='TNR',fontSize=11,leading=13,spaceAfter=7)
    title = ParagraphStyle('title',parent=normal,fontName='TNRB',fontSize=15,leading=17,spaceAfter=10)
    heading = ParagraphStyle('heading',parent=normal,fontName='TNRB',spaceBefore=8,spaceAfter=8,keepWithNext=True)
    small = ParagraphStyle('table',parent=normal,fontSize=9.5,leading=11,spaceAfter=0)
    para = lambda s,style=normal: Paragraph(escape(str(s)).replace('\n','<br/>'),style)
    flow=[para(p['sender']),Spacer(1,5),para('An: '+p['recipient']),para(f"{dd(p['date'])} | Zeichen: {p['ref']}"),para(p['title'],title),para(project)]
    for h,b in p['sections']: flow.extend([para(h,heading),para(b)])
    if p.get('table'):
        headers, rows=p['table']
        widths = [49,207,47,65,113] if len(headers)==5 else [70,170,95,146] if len(headers)==4 else [161]*3
        table=Table([[para(fmt(v),small) for v in row] for row in [headers]+rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7ECE9')),('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        flow += [Spacer(1,5),table,Spacer(1,8)]
    flow.append(KeepTogether([Spacer(1,5),para(p['signer'])]))
    def footer(canvas, doc):
        canvas.setFont('TNR',9); canvas.drawString(57,29,p['ref']); canvas.drawRightString(538,29,f'Seite {doc.page}')
    SimpleDocTemplate(str(path),pagesize=(595.28,841.89),leftMargin=57,rightMargin=57,topMargin=40,bottomMargin=46,
                      author='Klotzkette',title=p['title']).build(flow,onFirstPage=footer,onLaterPages=footer)


def print_profile(path):
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    with zipfile.ZipFile(path) as z: entries={i.filename:(i,z.read(i)) for i in z.infolist()}
    for name,(info,data) in list(entries.items()):
        if not name.startswith('xl/worksheets/sheet') or not name.endswith('.xml'): continue
        root=ET.fromstring(data)
        pr=root.find(f'{{{ns}}}sheetPr')
        if pr is None: pr=ET.Element(f'{{{ns}}}sheetPr');root.insert(0,pr)
        setuppr=ET.SubElement(pr,f'{{{ns}}}pageSetUpPr');setuppr.set('fitToPage','1')
        for tag in ['pageMargins','pageSetup']:
            old=root.find(f'{{{ns}}}{tag}')
            if old is not None: root.remove(old)
        ET.SubElement(root,f'{{{ns}}}pageMargins',dict(left='0.3',right='0.3',top='0.4',bottom='0.4',header='0.2',footer='0.2'))
        ET.SubElement(root,f'{{{ns}}}pageSetup',dict(paperSize='9',orientation='landscape',fitToWidth='1',fitToHeight='1'))
        entries[name]=(info,ET.tostring(root,encoding='utf-8',xml_declaration=True))
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for info,data in entries.values(): z.writestr(info,data)


def make_book(path,p):
    qa=QA_ROOT/path.parent.name/path.stem;qa.mkdir(parents=True,exist_ok=True)
    link=qa/'node_modules'
    if not link.exists(): link.symlink_to(RUNTIME/'node/node_modules',target_is_directory=True)
    staged=qa/path.name
    spec={**p,'path':str(staged),'qa':str(qa)}
    result=subprocess.run([str(NODE),'--input-type=module','-e',SHEET_JS],cwd=qa,input=json.dumps(spec),text=True,capture_output=True)
    if result.returncode: raise RuntimeError(result.stderr[-6000:]+result.stdout[-3000:])
    print_profile(staged)
    shutil.copyfile(staged,path)
    print(result.stdout.strip(),flush=True)


def make_mail(path,p):
    msg=EmailMessage(policy=policy.SMTP)
    msg['From']=p['sender'];msg['To']=p['recipient'];msg['Date']=format_datetime(datetime.fromisoformat(p['date']))
    msg['Subject']=p['title'];msg['Message-ID']=f"<{path.parent.name}.{path.stem}@aktenpost.example>"
    msg.set_content(p['body'])
    for name in p['attachments']:
        content=(path.parent/name).read_bytes()
        mime=mimetypes.guess_type(name)[0] or 'application/octet-stream'
        main,sub=mime.split('/',1)
        msg.add_attachment(content,maintype=main,subtype=sub,filename=name)
    path.write_bytes(msg.as_bytes())


def make_drawing(path,p):
    from PIL import Image,ImageDraw,ImageFont
    im=Image.new('RGB',(1600,1050),'white');draw=ImageDraw.Draw(im)
    font=ImageFont.truetype(str(FONTDIR/'Arial.ttf'),27)
    bold=ImageFont.truetype(str(FONTDIR/'Arial Bold.ttf'),34)
    draw.text((65,45),'Gerätehalle Wesertor | Baufeld 16.03.2026',fill='black',font=bold)
    draw.text((65,100),'Lageskizze der Tragwerksplanung, ohne Maßstab',fill='black',font=font)
    draw.rectangle((240,245,1310,740),outline='#202020',width=5)
    draw.rectangle((250,590,1300,730),fill='#F1D8CC',outline='#9E512B',width=4)
    draw.text((300,630),'Südtrasse gesperrt | Fundamenttaschen S3 / S4',fill='black',font=font)
    draw.text((310,370),'Nordtakt: Schalung und Bewehrung zugänglich',fill='black',font=font)
    for x in [300,600,900,1200]:
        for y in [280,550]: draw.rectangle((x-20,y-20,x+20,y+20),fill='#C6D8CE',outline='black',width=3)
    draw.line((120,460,230,460),fill='black',width=5)
    draw.polygon([(230,460),(208,446),(208,474)],fill='black')
    draw.text((65,410),'Zufahrt',font=font,fill='black')
    draw.line((1420,320,1420,190),fill='black',width=5)
    draw.polygon([(1420,175),(1404,200),(1436,200)],fill='black')
    draw.text((1406,130),'N',font=bold,fill='black')
    draw.text((65,820),'Freigabe Nord/West: F17B vom 10.03.2026',font=font,fill='black')
    draw.text((65,865),'F17C Süd am Zeichnungsdatum noch nicht freigegeben.',font=font,fill='black')
    draw.text((65,955),'Gez. Dr. Rolf Stein | Tragwerk Hunte | HW-26 / Skizze S06',font=font,fill='black')
    im.save(path,optimize=True)


def metadata(folder,c):
    import yaml
    inventory='\n'.join(f"| {i:02} | [{p['file']}]({p['file']}) | {dd(p['date'])} | {p['title']} |" for i,p in enumerate(c['pieces'],1))
    warning='> '+NOTICE.replace('\n\n','\n>\n> ')
    readme=f'''<!-- decimal-headings -->
# 1. {c['title']}

## 1.1. Gegenstand

{c['description']}

Plugin: `bauwirtschaft-fortgeschrittene`. Fallkennung: `{c['slug']}`.
Genau ein Ablauf aus Insert 2 des bereitgestellten Seminar-Inserts ist dieser Akte zugeordnet.
{len(c['pieces'])} eigenständige Originalunterlagen. Die zentralen drei Downloadformate werden durch den Paketbau ergänzt.

## 1.2. Herkunft

<!-- reserved-example-contacts -->

Personen, Unternehmen, Behördenorganisationen, Adressen und Projektvorgänge sind erfunden. Die Ortsnamen sind real. Auch Unterlagen öffentlicher Stellen sind eigens erstellte Fiktion und keine echten Bekanntmachungen oder Urkunden. Alle E-Mail-Kontakte verwenden reservierte `.example`-Domains. Es werden keine amtlichen Logos, Siegel, realen Bankdaten oder fingierten Fotografien verwendet. Eine Baufeldskizze ist eine gezeichnete technische Darstellung.

## 1.3. Originalunterlagen

{warning}

| Nr. | Datei | Datum | Inhalt |
| --- | --- | --- | --- |
{inventory}

## 1.4. Reproduktion und Grenzen

Erzeugung: `scripts/build-bau-rundum-vertiefung.py --case {c['slug']}`.
Prüfung: derselbe Aufruf mit `--check`; zusätzliche native Neuberechnung und Quellrendering mit `--qa`.
Das Datenskript `scripts/bau_rundum_vertiefung_daten.py` enthält ausschließlich die fünf zugewiesenen Abläufe.
Die interne `rubric.yaml` gehört nicht zu den Arbeitsunterlagen. Technische Prüfdateien liegen außerhalb des Repositorys.
Forderungen und Erklärungen sind Stimmen der Beteiligten, keine rechtlichen Ergebnisse. Es werden keine Live-Modelltests oder abgeschlossenen rechtlichen Quellenprüfungen behauptet. Die rechtliche Quellenprüfung bleibt Aufgabe der bearbeitenden Plugin-Agenten.

Autor: Klotzkette <39582916+Klotzkette@users.noreply.github.com>. Zugeordneter Pluginstand: 445.33.1. Akten-Begleitrelease: bauwirtschaft-rundum-v445.33.2.
'''
    if not (folder/'README.md').exists():
        (folder/'README.md').write_text(readme,encoding='utf-8')
    rubric=dict(name=c['slug'],plugin='bauwirtschaft-fortgeschrittene',description=c['description'],
                checks=[dict(id=f'sachpruefung-{i}',check_type='human_review',description=s) for i,s in enumerate(c['checks'],1)] +
                       [dict(id=f'original-{i:02}',check_type='file_exists',path=p['file'],description='Eigenständige Originalunterlage vorhanden.') for i,p in enumerate(c['pieces'],1)])
    if not (folder/'rubric.yaml').exists():
        (folder/'rubric.yaml').write_text(yaml.safe_dump(rubric,allow_unicode=True,sort_keys=False),encoding='utf-8')


def generate(c, new_only=False):
    folder=ROOT/'testakten'/c['slug'];folder.mkdir(parents=True,exist_ok=True)
    qa=QA_ROOT/c['slug'];qa.mkdir(parents=True,exist_ok=True)
    preserved={p['file']:hashlib.sha256((folder/p['file']).read_bytes()).hexdigest()
               for p in c['pieces'] if new_only and (folder/p['file']).is_file()}
    for sidecar in folder.glob('*.xlsx.inspect.ndjson'):
        shutil.move(str(sidecar),str(qa/sidecar.name))
    oldqa=folder/'.qa'
    if oldqa.exists():
        shutil.copytree(oldqa,qa,dirs_exist_ok=True,symlinks=True)
        shutil.rmtree(oldqa)
    for p in c['pieces']:
        path=folder/p['file'];ext=path.suffix
        if p['file'] in preserved:continue
        if ext=='.eml': continue
        if ext=='.docx': make_docx(path,p,c['project'])
        elif ext=='.pdf': make_pdf(path,p,c['project'])
        elif ext=='.csv':
            with path.open('w',encoding='utf-8',newline='') as f:
                w=csv.writer(f,delimiter=';');w.writerow(p['headers']);w.writerows(p['rows'])
        elif ext=='.xlsx': make_book(path,p)
        elif ext=='.png': make_drawing(path,p)
        elif ext=='.txt':
            text=p['text']
            if '{hashes}' in text:
                hashes='\n'.join(hashlib.sha256(f.read_bytes()).hexdigest()+'  '+f.name for f in sorted(folder.glob('0[567]_Angebot_*.pdf')))
                text=text.replace('{hashes}',hashes)
            path.write_text(text,encoding='utf-8')
        else: raise ValueError(ext)
    for p in c['pieces']:
        if p['file'].endswith('.eml') and p['file'] not in preserved:make_mail(folder/p['file'],p)
    for name,digest in preserved.items():
        assert hashlib.sha256((folder/name).read_bytes()).hexdigest()==digest,(folder,name)
    metadata(folder,c)
    return check(c)


def check(c):
    from openpyxl import load_workbook
    from docx import Document
    from pypdf import PdfReader
    from testakte_file_filter import include_in_working_dump
    folder=ROOT/'testakten'/c['slug']; counts=Counter();formula_count=0;pdf_pages=0;csv_rows=0;attachments=0
    originals=[folder/p['file'] for p in c['pieces']]
    assert len(originals)>=20 and len({p.suffix for p in originals})>=4
    additions=[p for p in originals if int(p.name.split('_',1)[0])>16]
    assert len(additions)>=4 and len({p.suffix for p in additions})>=2
    readme=(folder/'README.md').read_text(encoding='utf-8')
    assert all(f']({p.name})' in readme for p in additions)
    exported=[p for p in folder.rglob('*') if include_in_working_dump(p,folder)]
    assert set(exported)==set(originals), (set(exported)-set(originals),set(originals)-set(exported))
    for p,path in zip(c['pieces'],originals):
        assert path.is_file() and path.stat().st_size>0
        counts[path.suffix]+=1
        if path.suffix=='.eml':
            m=BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            for h in ['From','To','Date','Subject','Message-ID','MIME-Version']:assert m[h],(path,h)
            assert m.get_body(preferencelist=('plain',)) is not None
            assert m.get_body(preferencelist=('plain',)).get_content().replace('\r\n','\n').rstrip()==p['body'].rstrip()
            actual={a.get_filename():a.get_payload(decode=True) for a in m.iter_attachments()}
            assert set(actual)==set(p['attachments'])
            for name,data in actual.items():assert data==(folder/name).read_bytes();attachments+=1
            assert all(a.domain.endswith('.example') for h in ['From','To'] for a in m[h].addresses)
        elif path.suffix=='.csv':
            with path.open(encoding='utf-8',newline='') as f: rows=list(csv.reader(f,delimiter=';'))
            assert len(rows)>=5 and all(len(r)==len(rows[0]) for r in rows)
            assert rows[0]==p['headers']
            csv_rows+=len(rows)-1
        elif path.suffix=='.xlsx':
            wf=load_workbook(path,data_only=False);wv=load_workbook(path,data_only=True)
            for sheet in wf:
                for row in sheet:
                    for cell in row:
                        if cell.data_type=='f':
                            formula_count+=1;v=wv[sheet.title][cell.coordinate].value
                            assert v is not None and not str(v).startswith('#'),(path,cell.coordinate,v)
            for s in p['sheets']:
                for cell,expected in s['controls'].items():assert abs(wv[s['name']][cell].value-expected)<.011,(path,cell)
                for i,row in enumerate(s['rows']+s['tail'],6):
                    for j,value in enumerate(row,1):
                        if isinstance(value,str) and not value.startswith('='):
                            actual=wf[s['name']].cell(i,j).value
                            assert actual==value or (value=='' and actual is None),(path,i,j,value,actual)
        elif path.suffix=='.docx':
            d=Document(path);text='\n'.join(x.text for x in d.paragraphs)
            assert len(text)>900,(path,len(text))
            assert NOTICE.splitlines()[0] not in text
        elif path.suffix=='.pdf':
            pdf=PdfReader(path);pdf_pages+=len(pdf.pages)
            text='\n'.join(x.extract_text() for x in pdf.pages)
            assert len(text)>900,(path,len(text))
            assert NOTICE.splitlines()[0] not in text
    numeric = numeric_regression(c,folder)
    result=dict(case=c['slug'],originals=len(originals),types=dict(counts),pdf_pages=pdf_pages,csv_data_rows=csv_rows,
                formula_cells=formula_count,byte_identical_attachments=attachments,numeric_regression=numeric,
                new_originals=len(additions),addition_checks=addition_regression(c,folder),
                preserved_originals=len(originals)-len(additions))
    qa=QA_ROOT/c['slug'];qa.mkdir(parents=True,exist_ok=True)
    (qa/'checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False),flush=True)
    return result


def numeric_regression(c,folder):
    """Unabhängige Kontrollwerte aus gespeicherten Eingangsbelegen, kein Akteninhalt."""
    from openpyxl import load_workbook
    from pypdf import PdfReader
    def rows(name):
        with (folder/name).open(encoding='utf-8',newline='') as f:
            return list(csv.DictReader(f,delimiter=';'))
    def pdftext(name):
        return '\n'.join(p.extract_text() for p in PdfReader(folder/name).pages)
    D=Decimal
    slug=c['slug']
    if slug.endswith('lemgo'):
        payments=sum(D(r['Betrag_EUR']) for r in rows('06_Kreditorenbewegungen.csv') if r['Art']=='Auszahlung')
        assert payments==D('100639.08')
        q=rows('03_Aufmass_Vormonate.csv')
        prices=[D(x) for x in ['12500','185','1450','92','42','198','37','1745']]
        for key,expected in [('Mai_kumulativ','39400'),('Juni_kumulativ','92560')]:
            assert sum(D(r[key])*p for r,p in zip(q,prices))==D(expected)
        sheet=load_workbook(folder/'11_Abschlagsrechnung_3.xlsx',data_only=True)['WB26098']
        assert sheet['A6'].value=='01.010'
        assert sheet['B20'].number_format=='#,##0.00' and sheet['B21'].number_format=='#,##0.00'
        assert D(str(sheet['F22'].value)).quantize(D('.01'))==D('70657.11')
        assert '44.541,70' in pdftext('04_Abschlagsrechnung_1.pdf')
        assert '60.097,38' in pdftext('05_Abschlagsrechnung_2.pdf')
        return dict(payments_eur=str(payments),cumulative_net_eur='151522.50',requested_gross_eur='70657.11')
    if slug.endswith('hameln'):
        archive=rows('04_Referenzarchiv.csv')
        assert len({r['Kennung'] for r in archive})==6
        for r in archive:datetime.fromisoformat(r['Abschluss_Phase8'])
        sheet=load_workbook(folder/'10_Personaleinsatz.xlsx',data_only=True)['Kapazität']
        assert sum(sheet.cell(i,2).value for i in range(6,15))==328
        assert sum(sheet.cell(i,6).value for i in range(6,15))==-14
        assert sheet['F6'].value==-16 and sheet['F12'].value==-2
        return dict(archive_records=6,weekly_capacity_sum=328,weekly_balance_sum=-14)
    if slug.endswith('goslar'):
        raw=rows('08_Preisdaten_Export.csv');assert len(raw)==24
        controls={'Rautenbau':('332800','332800'),'Okerfenster':('325120','325120'),'Bergglas':('329880','329440')}
        for name,(calculated,stated) in controls.items():
            selected=[r for r in raw if r['Bieter']==name]
            assert len(selected)==8 and len({r['Pos'] for r in selected})==8
            assert sum(D(r['Menge'])*D(r['EP_EUR_netto']) for r in selected)==D(calculated)
            assert sum(D(r['GP_eingetragen_EUR']) for r in selected)==D(stated)
        assert D('329880')*D('.98')==D('323282.40')
        assert '329.440,00' in pdftext('04_Oeffnungsniederschrift.pdf')
        text=(folder/'15_Exportprotokoll.txt').read_text()
        for path in folder.glob('0[567]_Angebot_*.pdf'):assert hashlib.sha256(path.read_bytes()).hexdigest() in text
        return dict(calculated_bid_net_eur=['332800.00','325120.00','329880.00'],berg_stated_net_eur='329440.00',hashes=3)
    if slug.endswith('minden'):
        def weekdays(start,end):
            first=datetime.fromisoformat(start);last=datetime.fromisoformat(end)
            return {first+timedelta(days=i) for i in range((last-first).days+1) if (first+timedelta(days=i)).weekday()<5}
        first=weekdays('2026-03-09','2026-03-20');second=weekdays('2026-03-16','2026-03-27')
        assert (len(first),len(second),len(first&second),len(first|second))==(10,10,5,15)
        staff=rows('10_Personalstunden.csv')
        assert sum(D(r['Zusatzstunden_gesamt']) for r in staff)==D('144')
        actual=rows('03_Isttermine.csv')
        assert len(actual)==12 and actual[-1]['Istende']=='2026-06-08'
        assert (datetime.fromisoformat(actual[-1]['Istende'])-datetime(2026,5,15)).days==24
        assert next(r for r in actual if r['Vorgang']=='V03')['Istende']=='2026-03-30'
        return dict(disturbance_days=[10,10],overlap_weekdays=5,union_weekdays=15,extra_hours=144,handover_calendar_shift=24)
    if slug.endswith('northeim'):
        lengths=rows('06_Aufmass_Leitungsachsen.csv')
        assert sum(D(r['Länge_m']) for r in lengths)==D('280')
        assert sum(D(r['Länge_m']) for r in lengths if r['Abschnitt'] in {'S1','S2'})==D('136')
        assert D('280')*(D('128.97')-D('84.71'))+D('1850')+4*D('640')==D('16802.80')
        assert len(rows('14_Terminnotizen.csv'))==6
        return dict(total_length_m=280,south_length_m=136,old_ep_eur='84.71',new_ep_eur='128.97',offer_net_eur='16802.80')
    raise AssertionError(slug)


def addition_regression(c,folder):
    """Kontrolliert Rohdaten der Ergänzungen, ohne Ergebnisse in die Akten zu schreiben."""
    from docx import Document
    from pypdf import PdfReader
    from openpyxl import load_workbook
    def rows(name):
        with (folder/name).open(encoding='utf-8',newline='') as f:
            return list(csv.DictReader(f,delimiter=';'))
    def document(name):
        path=folder/name
        if path.suffix=='.docx':return '\n'.join(p.text for p in Document(path).paragraphs)
        return '\n'.join(p.extract_text() for p in PdfReader(path).pages)
    D=Decimal
    slug=c['slug']
    if slug.endswith('lemgo'):
        readings=rows('18_Zwischenzaehler_Sued.csv')
        values=[D(r['Stand_kWh']) for r in readings]
        assert len(values)==6 and all(b>=a for a,b in zip(values,values[1:]))
        assert values[-1]-values[0]==D('46.2') and values[2]-values[1]==D('10.5')
        assert '4.000,00 EUR' in document('20_Arbeitsbericht_Sockelnacharbeit.docx')
        return dict(meter_readings=6,shared_meter_kwh='46.2',weekend_interval_kwh='10.5')
    if slug.endswith('hameln'):
        roles=document('17_ARGE_Leistungsabgrenzung.docx')
        assert '40 Prozent' in roles and '60 Prozent' in roles
        assert '32 Wochenstunden' in document('20_Personalgespraech_Aydin.docx')
        assert '26 Wochenstunden' in document('18_Riedhof_Abstimmung_Januar.pdf')
        assert '1.180,00 EUR' in document('19_Versicherer_Risikofragen.pdf')
        return dict(arge_fee_shares=[40,60],aydin_contract_hours=32,riedhof_kruse_hours=26)
    if slug.endswith('goslar'):
        table=Document(folder/'17_Elementabmessungen_F1.docx').tables[0]
        areas=[D(r.cells[3].text.replace(',','.')) for r in table.rows[1:]]
        quantities=[int(r.cells[2].text) for r in table.rows[1:]]
        assert areas==[D('72'),D('72'),D('72'),D('24')]
        assert sum(areas)==D('240') and sum(quantities)==98
        journal=rows('18_Zugriffsjournal_Pruefraum.csv')
        assert len(journal)==7 and journal[0]['Ergebnis']=='gesperrt'
        assert datetime.fromisoformat(journal[1]['Zeitpunkt'])>datetime.fromisoformat('2026-10-05T10:00:00+02:00')
        return dict(element_groups=4,elements=98,total_area_m2=240,access_events=7)
    if slug.endswith('minden'):
        cards=rows('18_Geraetestunden_MK44.csv')
        assert sum(D(r['Betriebsstunden']) for r in cards)==D('19')
        invoice=document('17_Mietrechnung_Mobilkran.pdf')
        assert D('22')*D('420')==D('9240') and '9.240,00 EUR' in invoice and '10.995,60 EUR' in invoice
        hours={r['Datum']:D(r['Zusatzstunden_gesamt']) for r in rows('10_Personalstunden.csv') if r['Datum'].startswith('2026-04-')}
        forecast=D(str(load_workbook(folder/'13_Forderungsberechnung.xlsx',data_only=True)['BZ01']['D10'].value))
        recorded=sum(hours.values())
        daily_reduction=[max(hours.values())-hours[day] for day in ['2026-04-22','2026-04-28']]
        message=BytesParser(policy=policy.default).parsebytes((folder/'21_Fricke_Geraeteunterlagen.eml').read_bytes())
        assert 'jeweils acht Stunden weniger' in message.get_body(preferencelist=('plain',)).get_content()
        assert daily_reduction==[D('8'),D('8')] and forecast-sum(daily_reduction)==recorded==D('144')
        return dict(selected_cards=6,selected_operating_hours=19,crane_rent_net_eur=9240,
                    forecast_extra_hours=int(forecast),recorded_extra_hours=int(recorded))
    if slug.endswith('northeim'):
        stock=rows('18_Lagerbestaende_DN150.csv')
        sealed=sum(D(r['Länge_m']) for r in stock if r['Zustand']=='Bund geschlossen')
        assert sum(D(r['Länge_m']) for r in stock)==D('300') and sealed==D('240')
        assert sealed*D('34')*D('.9')==D('7344')
        assert '7.344,00 EUR' in document('19_Ruecknahmeauskunft_Rohrkontor.pdf')
        return dict(stock_length_m=300,sealed_length_m=240,conditional_return_net_eur=7344)
    raise AssertionError(slug)


def docx_renderer_qa(c, from_number=1):
    renderer=Path.home()/'.codex/plugins/cache/openai-primary-runtime/documents/26.905.11957/skills/documents/render_docx.py'
    for p in c['pieces']:
        if int(p['file'].split('_',1)[0])<from_number:continue
        if not p['file'].endswith('.docx'):continue
        path=ROOT/'testakten'/c['slug']/p['file']
        output=QA_ROOT/c['slug']/'docx-renderer'/path.stem
        result=subprocess.run([sys.executable,str(renderer),str(path),'--output_dir',str(output),'--dpi','110','--emit_pdf'],capture_output=True,text=True,timeout=120)
        assert result.returncode==0,(path,result.stderr[-2000:])
        assert list(output.glob('page-*.png')),path
    print(c['slug'],'DOCX-Skill-Renderer: bestanden',flush=True)


def normalize_docx_fonts(c):
    """Ändert nur konkurrierende Schriftattribute, nicht den Dokumenttext."""
    from docx import Document
    from docx.oxml.ns import qn
    folder=ROOT/'testakten'/c['slug']
    for p in c['pieces']:
        if not p['file'].endswith('.docx'):continue
        path=folder/p['file'];d=Document(path)
        for name in ['Normal','Title','Heading 1','Heading 2']:
            fonts=d.styles[name].element.get_or_add_rPr().rFonts
            for attr in list(fonts.attrib):
                if 'theme' in attr.lower():del fonts.attrib[attr]
            for script in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+script),'Times New Roman')
        d.save(path)
    for p in c['pieces']:
        if p['file'].endswith('.eml') and any(n.endswith('.docx') for n in p['attachments']):
            make_mail(folder/p['file'],p)


def visual_qa(c, from_number=1):
    from pdf2image import convert_from_path
    from pypdf import PdfReader
    from openpyxl import load_workbook
    from testakte_office_pdf import render_office_batch
    from PIL import Image,ImageDraw
    folder=ROOT/'testakten'/c['slug'];qa=QA_ROOT/c['slug'];qa.mkdir(parents=True,exist_ok=True)
    os.environ['SOFFICE']=str(SOFFICE)
    pieces=[p for p in c['pieces'] if int(p['file'].split('_',1)[0])>=from_number]
    paths=[folder/p['file'] for p in pieces if Path(p['file']).suffix in {'.docx','.xlsx'}]
    rendered=render_office_batch(paths)
    assert set(rendered)==set(paths),'Native Office-PDFs fehlen'
    for path,data in rendered.items(): (qa/(path.stem+'.pdf')).write_bytes(data)
    pages=[];pagecount=0
    for path in paths+[folder/p['file'] for p in pieces if p['file'].endswith('.pdf')]:
        pdfpath=qa/(path.stem+'.pdf') if path.suffix!='.pdf' else path
        pdf=PdfReader(pdfpath)
        for i in range(len(pdf.pages)):
            page=convert_from_path(str(pdfpath),dpi=95,first_page=i+1,last_page=i+1,poppler_path=str(RUNTIME/'bin/override'))[0]
            target=qa/f'{path.stem}-p{i+1}.png';page.save(target);pages.append(target);pagecount+=1
    mutations=0
    for p in pieces:
        if not p['file'].endswith('.xlsx'): continue
        native=qa/Path(p['file']).stem/'native';native.mkdir(exist_ok=True)
        source=qa/Path(p['file']).stem/'mutation.xlsx'
        profile=qa/Path(p['file']).stem/'office-profile'
        subprocess.run([str(SOFFICE),f'-env:UserInstallation={profile.as_uri()}','--headless','--convert-to','xlsx','--outdir',str(native),str(source)],check=True,capture_output=True,timeout=120)
        w=load_workbook(native/'mutation.xlsx',data_only=True)
        for s in p['sheets']:
            m=s['mutation'];assert abs(w[s['name']][m['result']].value-m['expected'])<.011
            mutations+=1
    for start in range(0,len(pages),9):
        contact=Image.new('RGB',(1200,1800),'#D8D8D8');draw=ImageDraw.Draw(contact)
        for j,path in enumerate(pages[start:start+9]):
            im=Image.open(path);im.thumbnail((385,550))
            x=(j%3)*400;y=(j//3)*600
            contact.paste(im,(x,y+30));draw.text((x+5,y+6),path.stem[:48],fill='black')
        contact.save(qa/f'kontakt-{start//9+1}.jpg',quality=85)
    (qa/'pages.json').write_text(json.dumps(dict(rendered_pages=pagecount,native_mutations=mutations,from_number=from_number,pages=[p.name for p in pages]),indent=2),encoding='utf-8')
    print(c['slug'], 'Office- und Original-PDF-Seiten:',pagecount,'native Mutationen:',mutations,flush=True)


def main():
    global QA_ROOT
    parser=argparse.ArgumentParser()
    parser.add_argument('--case',choices=list(CASES),action='append')
    parser.add_argument('--check',action='store_true')
    parser.add_argument('--qa',action='store_true')
    parser.add_argument('--docx-qa',action='store_true')
    parser.add_argument('--normalize-docx-fonts',action='store_true')
    parser.add_argument('--new-only',action='store_true',help='Erzeugt nur fehlende Originale; vorhandene Dateien bleiben bytegleich.')
    parser.add_argument('--qa-from',type=int,default=1,help='Kleinste laufende Nummer für visuelle QA.')
    parser.add_argument('--qa-root',type=Path,default=QA_ROOT)
    args=parser.parse_args()
    QA_ROOT=args.qa_root.resolve()
    if QA_ROOT==ROOT or ROOT in QA_ROOT.parents:parser.error('QA-Verzeichnis muss außerhalb des Repositorys liegen.')
    if args.new_only and args.normalize_docx_fonts:parser.error('Schriftnormalisierung widerspricht --new-only.')
    for key in args.case or CASES:
        c=CASES[key]
        if args.normalize_docx_fonts:normalize_docx_fonts(c)
        if args.check: check(c)
        else: generate(c,args.new_only)
        if args.qa:visual_qa(c,args.qa_from)
        if args.docx_qa:docx_renderer_qa(c,args.qa_from)


if __name__=='__main__': main()
