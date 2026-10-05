#!/usr/bin/env python3
"""Ergänzt zwölf Akten, ohne ein bisheriges natives Original zu überschreiben."""
import argparse,hashlib,importlib.util,json,re,shutil,subprocess,sys
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime,formataddr,getaddresses
from datetime import datetime
import mimetypes
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt
from kommunale_schriftverkehr_daten import cases,FOLDER,filename
from readme_decimal_headings import normalize_decimal_headings
from testakte_disclaimer import NOTICE_DE,NOTICE_EN
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('kommunal_helper',ROOT/'scripts/build-kommunale-haftpflicht-akten.py');helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)

def claim_body(case,item):
    body=item['body']
    if item['kind']=='claim' and 'Anlagenverzeichnis' not in body:
        headings=[int(x) for x in re.findall(r'^## (\d+)(?:\.|\s)',body,re.M)]
        number=max(headings,default=6)+1
        body+=f'\n\n## {number}. Anlagenverzeichnis\n\n'+'\n\n'.join(f"{x['label']}: {x['description']}\nDatei: {x['source']}" for x in case['exhibits'])
    return body

def safe_source(directory,name):
    path=directory/name
    if not path.resolve().is_relative_to(directory.resolve()) or path.is_symlink() or not path.is_file():
        raise ValueError('Fehlende/unsichere Aktenquelle: '+str(path))
    return path

def word(case,item,target):
    # Die ausformulierten Texte haben eigene Briefköpfe und Rubren.
    # Metadaten dürfen sie nicht als zweiten Kopf nochmals ausgeben.
    normalized={**item,'body':claim_body(case,item)}
    for key in ('sender','recipient','date'):
        normalized.pop(key,None)
    helper.word(normalized,target)
    doc=Document(target)
    doc.core_properties.author=item.get('sender','Aktenverwaltung')
    if item['kind']=='letter':
        title=next(p for p in doc.paragraphs if p.style.name=='Title')
        body=[p for p in doc.paragraphs if p is not title and p._p is not title._p]
        # Zwei vollständige Anschriften vor den Betreff stellen.
        for p in body[:2]:title._p.addprevious(p._p)
        if len(body)>2 and re.match(r'^\d{1,2}\.\d{1,2}\.\d{4}',body[2].text):
            title._p.addprevious(body[2]._p)
        else:
            date=doc.add_paragraph(item['date']);title._p.addprevious(date._p)
    helper.builder.native.stable_docx(doc,target)

def mail(case,item,target):
    msg=EmailMessage(policy=SMTP)
    for header,key,fallback in [('From','from','sender'),('To','to','recipient')]:
        addresses=getaddresses([item.get(key,item.get(fallback))])
        if not addresses or any('@' not in address for _,address in addresses):raise ValueError('Ungültiger Mailadressat')
        msg[header]=', '.join(formataddr(pair) for pair in addresses)
    msg['Subject']=item['title'];msg['Date']=format_datetime(datetime.fromisoformat(item['date']))
    msg['Message-ID']=f"<{case['slug']}.{target.stem}@akten.example>";msg.set_content(item['body'],charset='utf-8')
    directory=target.parent.parent
    for name in case.get('attachments',{}).get(item['file'],[]):
        source=target.parent/name
        if not source.is_file():source=directory/name
        if not source.resolve().is_relative_to(directory.resolve()) or source.is_symlink():raise ValueError('Unsichere MIME-Anlage')
        mime=mimetypes.guess_type(source.name)[0] or 'application/octet-stream';major,minor=mime.split('/',1)
        if major=='message':major,minor='application','octet-stream'
        msg.add_attachment(source.read_bytes(),maintype=major,subtype=minor,filename=source.name)
    if msg.is_multipart():msg.set_boundary('kommunal-'+case['slug']+'-'+target.stem)
    target.write_bytes(msg.as_bytes())

def metadata(case,directory):
    p=directory/'README.md';text=p.read_text();start='<!-- BEGIN kommunaler-schriftverkehr -->';end='<!-- END kommunaler-schriftverkehr -->'
    text=re.sub(re.escape(start)+r'.*?'+re.escape(end)+r'\n*','',text,flags=re.S)
    count=sum(1 for p in directory.iterdir() if p.is_file() and p.suffix in {'.eml','.docx','.txt','.xlsx','.pdf'})+6
    pdf=directory/'gesamt-pdf'/f"{case['slug']}_gesamt.pdf"
    pages=len(PdfReader(pdf).pages) if pdf.exists() else None
    text=re.sub(r'Das Gesamt-PDF umfasst \d+ Seiten\.',f'Das Gesamt-PDF umfasst {pages} Seiten.',text) if pages else text
    if pages:
        from testakte_einzelpdf_common import document_arcname_pairs
        pairs=document_arcname_pairs(directory)
        individual_pages=pages-sum(p.suffix in {'.docx','.xlsx','.pdf'} for p,a in pairs)
        text=re.sub(r'Das Einzel-PDF-ZIP enthält [^.\n]+ Unterlagen mit zusammen [^.\n]+ Seiten\.',f'Das Einzel-PDF-ZIP enthält {count} Unterlagen mit zusammen {individual_pages} Seiten.',text)
    # Historische Umfangsangaben bleiben als klar benannter Grundbestand erhalten.
    text=text.replace('Zwölf native Originalunterlagen.','Grundbestand: zwölf native Originalunterlagen.')
    text=text.replace('Die Akte enthält keine Musterlösung und keinen fertig ausgearbeiteten Vortrag.','Der Grundbestand bleibt ohne Musterlösung; der gesonderte Klageentwurf ist ein ausdrücklich bestelltes Arbeitsstück. Ein fertiger Vortrag wird nicht mitgeliefert.')
    text=text.replace('Die Aktenstücke enthalten keine Musterlösung.','Der Grundbestand bleibt ohne Musterlösung. Der zusätzlich bestellte Klageentwurf ist ein gesondertes Arbeitsstück.')
    text=text.replace('ein fertiger Vortrag oder eine Musterlösung ist nicht Bestandteil der Akte.','ein fertiger Vortrag ist nicht Bestandteil der Akte. Der zusätzlich bestellte Klagevorentwurf wird zeitlich und fachlich gesondert eingeordnet.')
    rows='\n'.join(f"| [{filename(d)}]({FOLDER}/{filename(d)}) | {d['title']} |" for d in case['documents'])
    notes=case.get('notes','');notes='\n\n'.join(notes) if isinstance(notes,list) else str(notes)
    notes=notes.replace('Anlagenverzeichnis aus exhibits ergänzen.','Das Anlagenverzeichnis ist im Word-Dokument enthalten.').replace('Das Anlagenverzeichnis wird aus exhibits angefügt','Das Word-Dokument enthält ein vollständiges Anlagenverzeichnis')
    section=f'''{start}

## Schriftverkehr und Klageentwurf

Zusätzlich zum unveränderten Grundbestand enthält die Akte drei E-Mails, zwei Briefe als PDF und einen ausgearbeiteten Klageentwurf als bearbeitbares Word-Dokument. Insgesamt stehen jetzt **{count} Originalunterlagen** zur Verfügung. Der Unterordner `{FOLDER}` trennt die Ergänzung vom bisherigen Akteneingang. In den flachen ZIPs wird der Unterordnername als Dateipräfix mitgeführt.

Der Klageentwurf ist ein bewusst zusätzlich bestelltes Arbeitsstück zur Überarbeitung oder Gegenprüfung. Er ersetzt nicht die offene Fallprüfung und ändert weder den ursprünglichen Auftraggeber noch den dort benannten Vertretungsumfang. Ein Entwurf aus Sicht der Gegenseite ist keine Übernahme ihres Mandats. Kein Schreiben und keine Klage wurden tatsächlich versandt oder eingereicht. Die Anlagen K1 ff. bezeichnen die konkreten, bereits vorhandenen Belege; dieselbe Datei wird dadurch kein zweiter Beweis oder weiterer Schaden.

{notes}

> {NOTICE_DE}
>
> {NOTICE_EN}

| Datei | Gegenstand |
| --- | --- |
{rows}

Die Briefe sind PDF-Originale dieser synthetischen Akte, keine Scans wirklicher Sendungen. Die Word-Klage enthält Anträge, Sachverhalt, Beweisantritte, rechtliche Begründung und Anlagenverzeichnis. Für das selbstständige Erarbeiten eines Entwurfs zunächst nur den Grundbestand öffnen; für Prüfung und Verbesserung anschließend den zusätzlichen Klageentwurf hinzunehmen.

{end}
'''
    # Vor Downloads einfügen, ohne die vorgeschriebene Hinweis-Link-Nachbarschaft zu trennen.
    match=re.search(r'(?:<!-- decimal-anchor --> <a id="downloads"></a>\n\n)?^## [^\n]*Downloads\s*$',text,re.M)
    if match:text=text[:match.start()]+section+'\n'+text[match.start():]
    else:text+='\n'+section
    p.write_text(normalize_decimal_headings(text))

def add_release_bookmarks(writer,directory):
    from testakte_einzelpdf_common import document_arcname_pairs
    from collections import Counter
    compact=lambda s:re.sub(r'\s+','',s)
    headers=[]
    for page in writer.pages:
        s=page.extract_text() or ''
        s=re.sub(r'^Akte:[^\n]*\nSeite \d+\n','',s)
        headers.append(compact(s[:400]))
    kickers={'.docx':'WORD-DOKUMENT (ORIGINAL-LAYOUT)','.xlsx':'EXCEL-TABELLE (BERECHNETE DRUCKFASSUNG)','.pdf':'PDF-ANHANG (ORIGINALDOKUMENT)','.eml':'E-MAILS','.txt':'NOTIZEN UND TEXTDATEIEN','.csv':'CSV-TABELLEN'}
    pairs=document_arcname_pairs(directory)
    stems=Counter(path.relative_to(directory).with_suffix('').as_posix() for path,arc in pairs)
    for path,arc in pairs:
        prefix=compact(kickers[path.suffix]+path.name)
        hits=[i for i,t in enumerate(headers) if t.startswith(prefix)]
        if len(hits)!=1:raise ValueError(f'Uneindeutiger tatsächlicher Dokumentkopf {path.name}: {hits}')
        page=hits[0]+(1 if path.suffix in {'.docx','.xlsx','.pdf'} else 0)
        stem=path.relative_to(directory).with_suffix('').as_posix()
        label=path.relative_to(directory).as_posix() if stems[stem]>1 else stem
        writer.add_outline_item(label,page)

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--qa-dir',type=Path,required=True);p.add_argument('--groups',nargs='*');p.add_argument('--cases',nargs='*');p.add_argument('--metadata-only',action='store_true');p.add_argument('--render-docx',type=Path);args=p.parse_args();args.qa_dir.mkdir(parents=True,exist_ok=True)
    selected=cases(args.groups)
    if args.cases:
        selected=[c for c in selected if c['slug'] in args.cases]
        if {c['slug'] for c in selected}!=set(args.cases):raise ValueError('Unbekannte Aktenauswahl')
    tasks=[];inventory=[]
    for case in selected:
        directory=ROOT/'testakten'/case['slug'];out=directory/FOLDER
        docs=case['documents'];assert len(docs)==6 and sum(d['kind']=='email' for d in docs)==3 and sum(d['kind']=='letter' for d in docs)==2 and sum(d['kind']=='claim' for d in docs)==1
        assert len({filename(d) for d in docs})==6 and all(Path(d['file']).name==d['file'] for d in docs)
        assert [e['label'] for e in case['exhibits']]==[f'K{i+1}' for i in range(len(case['exhibits']))]
        if not args.metadata_only:
            if not args.render_docx:raise ValueError('Der kanonische Word-Renderer ist für neue Briefe und Klagen erforderlich')
            out.mkdir(exist_ok=True)
            for item in docs:
                if item['kind']=='email':continue
                target=(args.qa_dir/'letter-sources'/case['slug']/item['file']) if item['kind']=='letter' else out/item['file'];target.parent.mkdir(parents=True,exist_ok=True)
                word(case,item,target);tasks.append((case,item,target))
        else:metadata(case,directory)
    if tasks:
        env=helper.builder.render_helpers.renderer_environment(args.qa_dir,Path(sys.executable).resolve().parents[3])
        def render(task):
            case,item,path=task;dest=args.qa_dir/'word'/case['slug']/path.stem;dest.mkdir(parents=True,exist_ok=True)
            for old_page in dest.glob('page-*.png'):old_page.unlink()
            run=subprocess.run([sys.executable,str(args.render_docx),str(path),'--output_dir',str(dest),'--emit_pdf','--dpi','100'],env=env,capture_output=True,text=True,timeout=300);(dest/'render.log').write_text(run.stdout+run.stderr)
            if run.returncode:raise RuntimeError(path.name+': '+run.stderr[-1500:])
            pdf=dest/(path.stem+'.pdf');images=sorted(dest.glob('page-*.png'));assert images and pdf.is_file()
            assert len(images)==len(PdfReader(pdf).pages) and {int(p.stem.split('-')[-1]) for p in images}==set(range(1,len(images)+1))
            if item['kind']=='letter':shutil.copyfile(pdf,ROOT/'testakten'/case['slug']/FOLDER/filename(item))
            return dict(case=case['slug'],kind=item['kind'],source=str(path),final=str(ROOT/'testakten'/case['slug']/FOLDER/filename(item)),pages=len(images),images=[str(p) for p in images],pdf=str(pdf))
        with ThreadPoolExecutor(max_workers=3) as pool:inventory=list(pool.map(render,tasks))
        for case in selected:
            directory=ROOT/'testakten'/case['slug'];out=directory/FOLDER
            pending=[d for d in case['documents'] if d['kind']=='email'];complete={filename(d) for d in case['documents'] if d['kind']!='email'}
            while pending:
                unresolved={(out/d['file']).resolve() for d in pending}
                def available(name):
                    source=out/name
                    if not source.is_file():source=directory/name
                    return source.is_file() and source.resolve().is_relative_to(directory.resolve()) and source.resolve() not in unresolved
                ready=[d for d in pending if all(available(n) for n in case.get('attachments',{}).get(d['file'],[]))]
                if not ready:raise ValueError('Fehlende oder zyklische MIME-Anlage '+case['slug'])
                for item in ready:mail(case,item,out/item['file']);pending.remove(item);complete.add(item['file'])
            for e in case['exhibits']:safe_source(directory,e['source'])
            metadata(case,directory)
        record=args.qa_dir/('word-render-'+('-'.join(args.groups) if args.groups else 'all')+'.json')
        if args.cases and record.exists():
            previous=json.loads(record.read_text());inventory=[r for r in previous if r['case'] not in args.cases]+inventory
        record.write_text(json.dumps(inventory,ensure_ascii=False,indent=2)+'\n')
    print('Akten:',len(selected),'gerenderte Word-Dokumente:',len(inventory))
if __name__=='__main__':main()
