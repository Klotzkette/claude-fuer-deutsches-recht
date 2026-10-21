#!/usr/bin/env python3
"""Ergänzt Weserfunken additiv; schützt alle 42 bisherigen Originale."""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import mimetypes
import os
import shutil
import subprocess
import tempfile
from datetime import datetime
from email import policy
from email.message import EmailMessage
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from akten_build_runtime import node_binary, serif_font_path

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / 'scripts/fixtures/fintech-bremen'
CASE = ROOT / 'testakten/fintech-darlehen-vertragsuebernahme-bremen'
DATA = json.loads((FIX / 'anfechtung-ergaenzung.json').read_text())
OUT = CASE / DATA['folder']
FONT = 'WeserfunkenErgaenzung'


def originals():
    return {str(p.relative_to(CASE)): hashlib.sha256(p.read_bytes()).hexdigest()
            for folder in ['01_eingang', '02_klageerwiderung', '03_korrespondenz']
            for p in sorted((CASE / folder).rglob('*'))
            if p.is_file() and OUT not in p.parents}


def verify_originals():
    expected = json.loads((FIX / 'anfechtung-original42-sha256.json').read_text())
    actual = originals()
    if actual != expected or len(actual) != 42:
        raise ValueError('Die 42 bisherigen Originaldateien weichen vom gesicherten Bestand ab')


def pdf(d):
    styles = {
        'body': ParagraphStyle('Body', fontName=FONT, fontSize=11, leading=14, spaceAfter=9),
        'small': ParagraphStyle('Small', fontName=FONT, fontSize=9, leading=11, spaceAfter=6),
        'title': ParagraphStyle('Title', fontName=FONT+'Bold', fontSize=15, leading=18, spaceAfter=14),
        'head': ParagraphStyle('Head', fontName=FONT+'Bold', fontSize=12, leading=15, spaceBefore=10, spaceAfter=9, keepWithNext=True)
    }
    p = lambda t, style='body': Paragraph(escape(str(t)).replace('\n','<br/>'), styles[style])
    story = [p(d['author'],'head'), p(d['address'],'small'), Spacer(1,10), p(d['recipient']), p(d['date']+' | '+d['reference'],'small'), p(d['title'],'title')]
    for section in d['sections']:
        story.append(p(section['heading'],'head'))
        story.extend(p(t) for t in section['paragraphs'])
        if section.get('table'):
            rows = [section['table']['headers'], *section['table']['rows']]
            t = Table([[p(c,'small') for c in row] for row in rows], colWidths=[119]*4, repeatRows=1)
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8ECEF')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
            story.extend([t,Spacer(1,12)])
        story.extend(p(t) for t in section.get('after',[]))
    def footer(canvas, doc):
        canvas.setFont(FONT,8)
        canvas.drawString(55,28,d['reference'])
        canvas.drawRightString(A4[0]-55,28,f'Seite {doc.page}')
    target = OUT/(d['id']+'.pdf')
    doc = SimpleDocTemplate(str(target),pagesize=A4,leftMargin=55,rightMargin=55,topMargin=38,bottomMargin=47,title=d['title'],author='Klotzkette',creator='Klotzkette',invariant=1)
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    if not all(p.extract_text().strip() for p in PdfReader(target).pages):
        raise ValueError('Leere PDF-Seite')


def word(d):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
    for style in doc.styles:
        if style.type == 1:
            style.font.name = 'Times New Roman'; style.font.size=Pt(11); style.font.color.rgb=RGBColor(0,0,0)
            style.paragraph_format.space_after=Pt(8); style.paragraph_format.line_spacing=1.1
            fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if key.endswith('Theme'): del fonts.attrib[key]
            for key in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+key),'Times New Roman')
            for border in style.element.xpath('./w:pPr/w:pBdr'): border.getparent().remove(border)
    doc.styles['Title'].font.size=Pt(15)
    doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.size=Pt(12)
    doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    doc.add_paragraph(d['author']+'\n'+d['address'])
    doc.add_paragraph(d['recipient'])
    doc.add_paragraph(d['date']+' | '+d['reference'])
    doc.add_paragraph(d['title'],'Title')
    for index,s in enumerate(d['sections']):
        heading=doc.add_paragraph(s['heading'],'Heading 1')
        # Keep the short concluding section with substantive context on page two.
        if index==1: heading.paragraph_format.page_break_before=True
        for text in s['paragraphs']: doc.add_paragraph(text)
    footer=sec.footer.paragraphs[0]
    footer.text=d['reference']+' | Seite '
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    for run in footer.runs: run.font.size=Pt(8)
    doc.core_properties.author=doc.core_properties.last_modified_by='Klotzkette'
    doc.core_properties.title=d['title']; doc.core_properties.language='de-DE'
    doc.core_properties.created=doc.core_properties.modified=datetime.fromisoformat(d['date'])
    doc.save(OUT/(d['id']+'.docx'))


def workbook_data():
    records=json.loads((FIX/'business-records.json').read_text())['documents']
    k7=next(d for d in records if d['id'].startswith('07_K7'))
    creditors=[]
    for i in (1,2):
        for row in k7['pages'][i]['table']['rows']:
            if row[0].startswith('Summe'): continue
            creditors.append({'name':row[0], 'invoice':row[1].split(' / ')[0], 'due':datetime.strptime(row[2],'%d.%m.%Y').date().isoformat(), 'amount':float(row[3].replace('.','').replace(',','.')), 'source':f'K7 Blatt {i+1}'})
    movements=[]
    with (CASE/'03_korrespondenz/30_Buchungsdaten_Geschaeftskonto.csv').open(newline='') as f:
        rows=list(csv.reader(f,delimiter=';'))
    for row in rows[1:]:
        if row[1]!='Vellio Finance GmbH': continue
        movements.append({'id':f'Z{len(movements)+1:02}', 'date':row[0], 'parties':'Weserfunken an Vellio', 'type':'Schuldnerzahlung', 'amount':-float(row[3]), 'status':'ausgeführt', 'source':row[5]+' / B3'})
    movements.extend([
      {'id':'E01','date':'2022-10-14','parties':'Vellio an Nexora','type':'Erwerbspreis','amount':500000,'status':'ausgeführt','source':'B4, 12:31 Uhr'},
      {'id':'A01','date':'2022-10-14','parties':'Nexora an Weserfunken','type':'Auszahlung','amount':500000,'status':'ausgeführt','source':'K3, 10:11 Uhr'},
      {'id':'V01','date':'2024-04-16','parties':'Nordhafen an Nordkai','type':'Vorschlag','amount':42840,'status':'nicht ausgeführt','source':'62 / 63 / 52'}
    ])
    assert len(creditors)==12 and sum(r['amount'] for r in creditors)==478380
    assert len(movements)==22
    return {'creditors':creditors,'movements':movements}


def mail(d):
    m=EmailMessage(policy=policy.SMTP)
    for k in ('from','to','cc'):
        if d.get(k): m[k.title()]=d[k]
    m['Date']=datetime.fromisoformat(d['date']); m['Subject']=d['subject']
    m['Message-ID']=f"<{d['id']}@{d['from'].split('@')[-1].rstrip('>')}>"
    m.set_content(d['body'])
    for name in d['attachments']:
        content=(OUT/name).read_bytes()
        # application/octet-stream preserves the original EML as a downloadable file.
        mime=mimetypes.guess_type(name)[0] if not name.endswith('.eml') else 'application/octet-stream'
        main,sub=(mime or 'application/octet-stream').split('/',1)
        m.add_attachment(content,maintype=main,subtype=sub,filename=name)
    if d['attachments']: m.set_boundary('=_weserfunken_'+d['id'])
    (OUT/(d['id']+'.eml')).write_bytes(bytes(m))


def inventory():
    path=CASE/'README.md'; text=path.read_text()
    text='\n'.join(line for line in text.splitlines() if not line.startswith('| ['+DATA['folder']+'/'))+'\n'
    rows=[]
    for key,ext in [('documents',None),('workbooks','xlsx'),('emails','eml')]:
        for item in DATA[key]:
            name=item['id']+'.'+(ext or item['format'])
            link=DATA['folder']+'/'+name
            rows.append(f"| [{link}]({link}) | {item.get('title',item.get('subject',''))} |")
    text=text.replace('<!-- END fintech-inventory -->','\n'.join(rows)+'\n\n<!-- END fintech-inventory -->')
    text=re.sub(r"\|\n(?:[ \t]*\n)+(?=\|)", "|\n", text)
    path.write_text(text)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--qa-dir',type=Path,required=True); args=parser.parse_args()
    verify_originals(); OUT.mkdir(parents=True,exist_ok=True); args.qa_dir.mkdir(parents=True,exist_ok=True)
    for suffix,bold in [('',False),('Bold',True)]: pdfmetrics.registerFont(TTFont(FONT+suffix,str(serif_font_path(bold))))
    for d in DATA['documents']: (pdf if d['format']=='pdf' else word)(d)
    modules=Path(os.environ.get('AKTEN_NODE_MODULES','')).expanduser()
    if not str(os.environ.get('AKTEN_NODE_MODULES','')) or not modules.is_dir(): raise RuntimeError('AKTEN_NODE_MODULES muss auf den bereitgestellten Artifact-tool-Modulordner zeigen')
    with tempfile.TemporaryDirectory(prefix='weserfunken-ergaenzung-') as temporary:
        t=Path(temporary); (t/'node_modules').symlink_to(modules,target_is_directory=True)
        shutil.copyfile(FIX/'anfechtung-workbooks.mjs',t/'build.mjs')
        (t/'input.json').write_text(json.dumps(workbook_data(),ensure_ascii=False))
        subprocess.run([node_binary(),str(t/'build.mjs'),str(t/'input.json'),str(OUT),str(args.qa_dir)],check=True)
    spec=importlib.util.spec_from_file_location('fintech_base_builder',ROOT/'scripts/build-fintech-bremen-akte.py')
    base=importlib.util.module_from_spec(spec); spec.loader.exec_module(base)
    for item in DATA['workbooks']: base.workbook_print_settings(OUT/(item['id']+'.xlsx'),item)
    for d in DATA['emails']: mail(d)
    inventory(); verify_originals()
    files=list(OUT.iterdir())
    if len(files)!=20: raise ValueError(f'20 Originaldateien erwartet, gefunden: {len(files)}')
    print('20 Ergänzungen erzeugt; 42 bestehende Originaldateien bytegleich.')


if __name__=='__main__': main()
