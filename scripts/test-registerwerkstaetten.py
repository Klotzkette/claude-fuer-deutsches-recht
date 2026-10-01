#!/usr/bin/env python3
"""Fachliche Grenzen, native Falldateien und Publikationsumfang der Registerpakete."""
import hashlib,json,re,sys,unittest
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from registerakten_transparenz import CASES as A
from registerakten_handelsregister import CASES as B
from registerakten_grundbuch import CASES as C
from registerakten_marken import CASES as D
from testakte_disclaimer import NOTICE_DE,NOTICE_EN
ROOT=Path(__file__).resolve().parents[1]
PLUGINS=['transparenzregister-assistent','handelsregister-assistent','grundbuchamt-assistent','markenamt-assistent']
class RegisterWerkstaetten(unittest.TestCase):
 def test_exact_skill_count_and_standalone_prompts(self):
  for slug in PLUGINS:
   with self.subTest(plugin=slug):
    folder=ROOT/slug;self.assertEqual(len(list((folder/'skills').glob('*/SKILL.md'))),11)
    for kind in ['werkstatt','schnellstart','hauptproblem']:
     p=folder/f'{slug}-{kind}.md';data=p.read_bytes();self.assertEqual(data,p.with_suffix('.txt').read_bytes())
     if kind!='werkstatt':self.assertLessEqual(len(data),7500);self.assertLessEqual(len(data.decode()),7500)
     else:self.assertGreater(len(data),50000)
    self.assertFalse((ROOT/'testakten/megaprompts'/f'{slug}.md').exists())
 def test_cases_have_real_originals_and_complete_pdfs(self):
  for group in [A,B,C,D]:
   self.assertEqual(len(group),3)
   for c in group:
    with self.subTest(case=c['slug']):
     folder=ROOT/'testakten'/c['slug'];self.assertGreaterEqual(len(c['documents']),20)
     formats={Path(d['file']).suffix for d in c['documents']};self.assertTrue({'.docx','.pdf','.xlsx','.eml','.txt'}<=formats)
     readme=(folder/'README.md').read_text();self.assertIn(NOTICE_DE,readme);self.assertIn(NOTICE_EN,readme)
     for d in c['documents']:
      p=folder/d['file'];self.assertTrue(p.is_file())
      if p.suffix=='.docx':
       doc=Document(p);self.assertEqual(doc.styles['Normal'].font.name,'Times New Roman');self.assertEqual(doc.styles['Normal'].font.size.pt,11)
       self.assertGreater(sum(len(t.text) for t in doc.paragraphs),250)
     reader=PdfReader(folder/'gesamt-pdf'/f"{c['slug']}_gesamt.pdf");text='\n'.join(p.extract_text() or '' for p in reader.pages)
     self.assertGreater(len(reader.pages),20);self.assertNotIn(NOTICE_DE,text);self.assertNotIn(NOTICE_EN,text)
 def test_spreadsheet_formula_results_are_cached_and_dates_formatted(self):
  for c in A+B+C+D:
   for d in c['documents']:
    if not d['file'].endswith('.xlsx'):continue
    file=ROOT/'testakten'/c['slug']/d['file'];w=load_workbook(file,data_only=True);f=load_workbook(file,data_only=False)
    for spec in d['sheets']:
     for cell,value in spec.get('expected',{}).items():
      with self.subTest(file=file.name,sheet=spec['name'],cell=cell):self.assertAlmostEqual(w[spec['name']][cell].value,value,places=7)
     for ri,row in enumerate(spec['rows'],2):
      for ci,val in enumerate(row,1):
       cell=f[spec['name']].cell(ri,ci)
       if isinstance(val,str) and re.fullmatch(r'\d{2}\.\d{2}\.\d{4}',val):self.assertNotEqual(cell.number_format,'General')
    w.close();f.close()
 def test_trademark_decision_identity_and_boundaries(self):
  for kind in ['werkstatt','schnellstart','hauptproblem']:
   t=(ROOT/'markenamt-assistent'/f'markenamt-assistent-{kind}.md').read_text()
   for ref in ['I ZB 58/25','C-337/22 P','T-545/24','T-64/25','C-412/24']:self.assertIn(ref,t)
   self.assertIn('17.09.2026',t);self.assertIn('05.02.2026',t)
   self.assertNotIn('Urteil vom 17.09.2026',t)
   self.assertIn('Eintragung',t);self.assertIn('Anmeldung',t);self.assertIn('Freigabe',t)
 def test_foreign_authority_and_lost_letter_have_actual_evidence(self):
  hr=' '.join(str(d) for c in B for d in c['documents']);self.assertIn('Sri Lanka',hr)
  gb=' '.join(str(d) for c in C for d in c['documents']);self.assertIn('Kopierer',gb);self.assertIn('Ersatzbrief',gb)
  tr=' '.join(str(d) for c in A for d in c['documents']);self.assertIn('Aino',tr)
 def test_review_hashes_track_actual_prompts(self):
  for slug in PLUGINS:
   profile=json.loads((ROOT/'quality/evals'/f'{slug}.json').read_text())
   for kind,key in [('werkstatt','workshop_review'),('schnellstart','mini_review')]:self.assertEqual(profile[key]['sha256'],hashlib.sha256((ROOT/slug/f'{slug}-{kind}.md').read_bytes()).hexdigest())
if __name__=='__main__':unittest.main()
