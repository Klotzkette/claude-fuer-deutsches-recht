#!/usr/bin/env python3
"""Local structural/content regressions, not legal or live-model certification."""

import argparse
import contextlib
import copy
import csv
import hashlib
import importlib.util
import io
import json
import re
import sys
import tempfile
import unittest
from datetime import datetime
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
from pathlib import Path

import yaml
from docx import Document
from pypdf import PdfReader

import quality_lab

ROOT = Path(__file__).resolve().parents[1]
SLUG = "ki-verordnung-hochrisiko-pruefer"
PLUGIN = ROOT / SLUG
CASE = ROOT / "testakten/ki-hochrisiko-bewerbungsauswahl-kassel"
FIXTURE = ROOT / "scripts/fixtures" / SLUG / "case.json"
PROFILE = ROOT / "quality/evals" / f"{SLUG}.json"
WARNING = (
    "Diese Testakte wurde mit KI generiert und ist ein Experiment. "
    "Benutzung auf eigene Verantwortung und eigene Gefahr.\n\n"
    "This test case file was generated with AI and is an experiment. "
    "Use at your own responsibility and risk."
)
QA_DIR = None


def normalized(text):
    return " ".join(text.split())


def plugin_link(source, link, root=PLUGIN, allow_directory=False):
    if re.match(r"https?://|mailto:|#", link):
        return None
    target = (source.parent / link.split("#", 1)[0]).resolve()
    exists = target.is_file() or (allow_directory and target.is_dir())
    if not target.is_relative_to(root.resolve()) or not exists:
        raise ValueError(f"Nonportable or missing link: {source.name}: {link}")
    return target


def load_builder():
    spec = importlib.util.spec_from_file_location("kassel_builder", ROOT / "scripts" / f"build-{SLUG}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class PluginTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = json.loads(FIXTURE.read_text(encoding="utf-8"))
        cls.records = cls.data["documents"]
        cls.profile = json.loads(PROFILE.read_text(encoding="utf-8"))

    def test_manifest_and_nine_substantive_skills(self):
        manifest = json.loads((PLUGIN / ".claude-plugin/plugin.json").read_text())
        self.assertEqual(manifest["name"], SLUG)
        self.assertRegex(manifest["version"], r"^\d+\.\d+\.\d+$")
        self.assertLessEqual(len(manifest["description"]), 300)
        self.assertIn("Anhang-III", manifest["description"])
        self.assertEqual(manifest["author"], {"name": "Klotzkette", "email": "39582916+Klotzkette@users.noreply.github.com"})
        skills = list((PLUGIN / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 9)
        for path in skills:
            text = path.read_text(encoding="utf-8")
            meta = yaml.safe_load(text.split("---", 2)[1])
            self.assertEqual(set(meta), {"name", "description"}, path)
            self.assertEqual(meta["name"], path.parent.name)
            self.assertRegex(meta["name"], r"^[a-z0-9-]{1,64}$")
            self.assertTrue(80 <= len(meta["description"]) <= 1024)
            self.assertGreater(len(text.split()), 240, path)
            for number in range(1, 7):
                self.assertRegex(text, rf"(?m)^## {number} [^\n]+\n\n")
            self.assertIn("11 pt", text)
            self.assertRegex(text, r"[Aa]usformuliert")

    def test_portable_links_in_every_plugin_markdown(self):
        for path in PLUGIN.rglob("*.md"):
            for link in re.findall(r"\]\(([^)]+)\)", path.read_text(encoding="utf-8")):
                # The repository README may navigate to the catalog; runtime instructions may not.
                navigation = path == PLUGIN / "README.md"
                plugin_link(path, link, root=ROOT if navigation else PLUGIN, allow_directory=navigation)

    def test_portability_guard_rejects_root_dependency(self):
        skill = PLUGIN / "skills/artikel-6-software-einstufen/SKILL.md"
        with self.assertRaises(ValueError):
            plugin_link(skill, "../../../references/zitierweise.md")
        with self.assertRaises(ValueError):
            plugin_link(skill, "../../references/missing.md")
        self.assertEqual(plugin_link(skill, "../../references/zitierweise.md"), PLUGIN / "references/zitierweise.md")

    def test_prompts_and_editorial_hashes(self):
        mini = (PLUGIN / f"{SLUG}-schnellstart.md").read_bytes()
        workshop = (PLUGIN / f"{SLUG}-werkstatt.md").read_bytes()
        self.assertLessEqual(len(mini), 7500)
        self.assertLessEqual(len(mini.decode()), 7500)
        self.assertGreater(len(workshop), 14000)
        for data, key in ((mini, "mini_review"), (workshop, "workshop_review")):
            self.assertEqual(hashlib.sha256(data).hexdigest(), self.profile[key]["sha256"])
        quality_lab.validate_profile(self.profile, SLUG, PLUGIN, ROOT)

    def test_profile_rejects_stale_hash_and_answer_key_input(self):
        changed = copy.deepcopy(self.profile)
        changed["mini_review"]["sha256"] = "0" * 64
        with self.assertRaises(quality_lab.LabError):
            quality_lab.validate_profile(changed, SLUG, PLUGIN, ROOT)
        changed = copy.deepcopy(self.profile)
        changed["cases"][0]["input_files"] = [str(PROFILE.relative_to(ROOT))]
        with self.assertRaises(quality_lab.LabError):
            quality_lab.validate_profile(changed, SLUG, PLUGIN, ROOT)

    def test_classification_reference_and_negative_control(self):
        reference = (PLUGIN / "references/artikel-6-und-anhang-iii.md").read_text()
        for number in range(1, 9):
            self.assertRegex(reference, rf"(?m)^### 2\.{number} ")
        for phrase in ("Absätze 1a bis 1c", "Abschnitt A und B", "Artikel 2 Absatz 2", "Profiling", "Artikel 43", "Artikel 49", "Artikel 73"):
            self.assertIn(phrase, reference)
        case = next(c for c in self.profile["cases"] if c["id"] == "zeugnisrechtsberatung-negativkontrolle")
        criteria = " ".join(c["text"] for c in case["criteria"])
        self.assertIn("Verneint eine automatische Zuordnung", criteria)
        self.assertIn("automatische", criteria)
        self.assertIn("Nummer 4 Buchstabe b", criteria)
        self.assertIn("Arbeitnehmerkanzlei", case["request"])

    def test_current_law_safeguards_remain_in_materials(self):
        reference = (PLUGIN / "references/rechtsstand-und-quellen.md").read_text()
        for phrase in ("2026/1744", "27. Juli 2026", "2. Dezember 2027", "2. August 2028", "Abschnitte 1 bis 3", "Artikel 6 Absatz 5", "vereinfacht, nicht abgeschafft", "Artikel 4a ist keine Deepfake-Ausnahme", "KI-Marktüberwachungs", "weiterhin Entwurf"):
            self.assertIn(phrase, reference)
        for path in PLUGIN.rglob("*.md"):
            self.assertNotIn(chr(167), path.read_text(), path)

    def test_exact_ten_native_source_records(self):
        expected = {record["file"] for record in self.records}
        actual = {p.name for p in CASE.iterdir() if re.match(r"\d{2}_", p.name)}
        self.assertEqual(len(expected), 10)
        self.assertEqual(expected, actual)
        self.assertEqual({Path(name).suffix for name in expected}, {".eml", ".txt", ".csv", ".pdf", ".docx"})
        self.assertEqual(self.data["employees"], 112)
        for record in self.records:
            text = " ".join(record.get("paragraphs", []))
            text += " ".join(p for page in record.get("pages", []) for p in page["paragraphs"])
            if not record["file"].endswith(".csv"):
                self.assertGreater(len(text.split()), 180, record["file"])
            self.assertNotRegex(text, r"Musterlösung|Erwartungshorizont|Sollantwort|Lösungsschlüssel")
            self.assertNotIn("Diese Testakte", text)

    def test_eml_headers_bodies_contacts_and_no_lost_attachments(self):
        for record in self.records:
            if not record["file"].endswith(".eml"):
                continue
            message = BytesParser(policy=policy.default).parsebytes((CASE / record["file"]).read_bytes())
            self.assertFalse(message.defects)
            self.assertEqual(normalized(str(message["Subject"])), normalized(record["subject"]))
            self.assertEqual(parsedate_to_datetime(message["Date"]), datetime.fromisoformat(record["date"]))
            self.assertEqual(normalized(message.get_body().get_content()), normalized("\n\n".join(record["paragraphs"])))
            self.assertEqual(list(message.iter_attachments()), [])
            for _, address in getaddresses([str(message["From"]), str(message["To"])]):
                self.assertTrue(address.endswith(".example"), address)

    def test_csv_exact_export_and_event_order(self):
        record = next(r for r in self.records if r["file"].endswith(".csv"))
        with (CASE / record["file"]).open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream, delimiter=";"))
        self.assertEqual(len(rows), 12)
        self.assertEqual({row["Bewerbung_ID"] for row in rows}, {f"A{n:02d}" for n in range(1, 13)})
        self.assertEqual(sorted(int(row["Rang"]) for row in rows), list(range(1, 13)))
        self.assertEqual([[row[column] for column in record["columns"]] for row in rows], [[str(v) for v in row] for row in record["rows"]])
        invited = {row["Bewerbung_ID"] for row in rows if row["Statuszeit"]}
        self.assertEqual(invited, {"A01", "A03", "A05", "A08"})
        self.assertEqual(sum(row["Status_am_20260917"] == "offen" for row in rows), 8)
        for row in rows:
            self.assertEqual(row["Vorgang"], self.data["reference"])
            self.assertTrue(0 <= int(row["Passungswert"]) <= 100)
            if row["Statuszeit"]:
                self.assertLess(datetime.fromisoformat(row["Bewertung_am"]), datetime.fromisoformat(row["Personenansicht_geoeffnet_am"]))
                self.assertLess(datetime.fromisoformat(row["Personenansicht_geoeffnet_am"]), datetime.fromisoformat(row["Statuszeit"]))
            self.assertTrue(all(not value.startswith(("=", "+", "@")) for value in row.values()))
        a07 = next(row for row in rows if row["Bewerbung_ID"] == "A07")
        self.assertEqual(a07["Rang"], "10")
        self.assertEqual(a07["Personenansicht_geoeffnet_am"], "")
        self.assertIn("Textschicht fehlt", a07["Dateihinweis"])

    def test_docx_complete_text_native_fonts_and_no_title_border(self):
        for record in self.records:
            if not record["file"].endswith(".docx"):
                continue
            document = Document(CASE / record["file"])
            texts = [p.text for p in document.paragraphs]
            self.assertEqual(document.core_properties.author, record["author"])
            self.assertEqual(texts[0], record["title"])
            for page in record["pages"]:
                self.assertIn(page["heading"], texts)
                for paragraph in page["paragraphs"]:
                    self.assertIn(paragraph, texts)
            for style_name in ("Normal", "Title", "Heading 1"):
                style = document.styles[style_name]
                self.assertEqual(style.font.name, "Times New Roman")
                self.assertFalse(style.element.xpath(".//w:pBdr"))
                fonts = style.element.get_or_add_rPr().rFonts
                self.assertFalse(any("theme" in key.lower() for key in fonts.attrib))
            self.assertEqual(document.styles["Normal"].font.size.pt, 11)
            self.assertEqual(len(document.element.xpath('.//w:br[@w:type="page"]')), 1)
            self.assertTrue(document.sections[0].footer._element.xpath(".//w:fldSimple"))

    def test_pdf_every_source_paragraph_and_page_label(self):
        record = next(r for r in self.records if r["file"].endswith(".pdf"))
        pdf = PdfReader(CASE / record["file"])
        self.assertEqual(len(pdf.pages), 2)
        for index, expected in enumerate(record["pages"]):
            text = normalized(pdf.pages[index].extract_text())
            self.assertIn(expected["heading"], text)
            self.assertIn(f"Seite {index + 1} von 2", text)
            self.assertIn(self.data["reference"], text)
            for paragraph in expected["paragraphs"]:
                self.assertIn(normalized(paragraph), text)
            self.assertNotIn("Diese Testakte", text)

    def test_readme_warnings_neutral_rubric_and_case_links(self):
        readme = (CASE / "README.md").read_text()
        self.assertIn(WARNING, readme)
        self.assertLess(readme.index(WARNING), readme.index("| Datei |"))
        self.assertIn(WARNING, (PLUGIN / "README.md").read_text())
        for link in re.findall(r"\]\(([^)]+)\)", readme):
            plugin_link(CASE / "README.md", link, root=ROOT, allow_directory=True)
        rubric = yaml.safe_load((CASE / "rubric.yaml").read_text())
        self.assertEqual(rubric["plugin"], SLUG)
        for check in rubric["checks"]:
            self.assertIn(check["check_type"], ("file_exists", "human_review"))
            if check["check_type"] == "file_exists":
                self.assertTrue((CASE / check["path"]).is_file())

    def test_builder_is_topic_scoped_and_reproduces_source_text(self):
        builder = load_builder()
        self.assertEqual(builder.CASE, CASE)
        with tempfile.TemporaryDirectory(prefix="kassel-regression-") as temp:
            with contextlib.redirect_stdout(io.StringIO()):
                builder.build(Path(temp))
            self.assertEqual({p.name for p in Path(temp).iterdir()}, {r["file"] for r in self.records})
            for record in self.records:
                output = Path(temp) / record["file"]
                original = CASE / record["file"]
                if output.suffix in (".eml", ".csv", ".txt"):
                    self.assertEqual(output.read_bytes(), original.read_bytes())
                elif output.suffix == ".docx":
                    self.assertEqual([p.text for p in Document(output).paragraphs], [p.text for p in Document(original).paragraphs])
                else:
                    self.assertEqual([p.extract_text() for p in PdfReader(output).pages], [p.extract_text() for p in PdfReader(original).pages])

    def test_rendered_docx_page_counts_and_no_paragraph_loss(self):
        if QA_DIR is None:
            self.skipTest("Use --qa-dir after native Office render; this check does not replace image inspection")
        for record in self.records:
            if not record["file"].endswith(".docx"):
                continue
            directory = QA_DIR / record["file"][:2]
            rendered = directory / (Path(record["file"]).stem + ".pdf")
            pdf = PdfReader(rendered)
            self.assertEqual(len(pdf.pages), len(record["pages"]))
            for index, expected in enumerate(record["pages"]):
                self.assertTrue((directory / f"page-{index + 1}.png").is_file())
                text = normalized(pdf.pages[index].extract_text())
                self.assertIn(f"Seite {index + 1} von 2", text)
                for paragraph in expected["paragraphs"]:
                    self.assertIn(normalized(paragraph), text)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--qa-dir", type=Path, help="Existing native-render directory containing 02/ and 10/")
    args, remaining = parser.parse_known_args()
    QA_DIR = args.qa_dir
    unittest.main(argv=[sys.argv[0], *remaining])
