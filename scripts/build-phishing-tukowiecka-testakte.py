#!/usr/bin/env python3
"""Fiktive deutsche Arbeitsakte zum Kernablauf C-70/25 reproduzierbar erzeugen.

Native DOCX/PDF/EML und zwei mit Artifact Tool berechnete XLSX. Rechtliche
Entwurfsabsätze bleiben in der kanonischen Datenquelle kontrollierbar.
"""
from pathlib import Path
import argparse, datetime as dt, hashlib, json, os, re, subprocess
from email.message import EmailMessage
from email.policy import SMTP
from email.parser import BytesParser
from email.utils import format_datetime
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
SLUG = 'phishing-kleinanzeigen-rabenstein-halle'
OUT = ROOT / 'testakten' / SLUG
QA = Path('/tmp/phishing-rabenstein-qa')
RUNTIME = Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies'
NOTICE = '> Diese Testakte wurde mit KI generiert und ist ein Experiment. Benutzung auf eigene Verantwortung und eigene Gefahr.\n>\n> This test case file was generated with AI and is an experiment. Use at your own responsibility and risk.'
DOCS = []
def add(n, ext, title, text):
    DOCS.append(dict(file=f'{n:02d}_{title.replace(" ", "_")}.{ext}', title=title.replace('_',' ').replace('Gedaechtnis','Gedächtnis').replace('Geraet','Gerät').replace('Rueckruf','Rückruf'), paragraphs=[p.strip() for p in text.strip().split('\n\n') if p.strip()]))

def contents():
    add(1,'docx','Mandatsauftrag_16_Oktober', '''Kanzlei Fenchel. Rechtsanwältin Daria Fenchel, Fiktiver Birkenbogen 8, 06108 Halle (Saale). Akte DF 166/26. Aufnahme am Freitag, 16.10.2026, 09.10 Uhr.

1. Auftrag und Beteiligte

Kunigunde Rabenstein, Fiktive Lindenhofgasse 17, 06108 Halle (Saale), ist private Kontoinhaberin bei der Saaleufer Bank AG, Fiktiver Uferplatz 4, 06108 Halle (Saale). Das Konto mit der internen Kundenkennung KR-4821 wird für ihren Haushalt geführt. Frau Rabenstein arbeitet als Bibliotheksassistentin. Sie hat am 12.10.2026 einen gebrauchten Lesesessel auf der Plattform Hofmarkt angeboten. Eine fremde Person hat anschließend 700,00 EUR von ihrem Konto überwiesen. Frau Rabenstein will diesen Betrag zurück und eine verständliche Erklärung der Bank.

Sie beauftragt die Kanzlei zunächst, die vorhandenen Belege zu ordnen und das Erstattungsschreiben zu überarbeiten. Ein Klageentwurf soll vorbereitet, aber noch nicht eingereicht werden. Sie wünscht keine Strafanzeige gegen Beschäftigte der Bank und keine Veröffentlichung ihres Falls. Die am 13.10.2026 an die Polizei übermittelte Sachverhaltsschilderung liegt als private Kopie vor. Eine amtliche Eingangsbestätigung fehlt noch. Ein anwaltliches Honorar ist gesondert zu vereinbaren; aus der Buchungsliste ergibt sich keine Kostenfreigabe.

2. Sachstand

Am Montagabend gelangte Frau Rabenstein über eine angebliche Kaufbestätigung auf eine Seite, die für sie wie der Zugang ihrer Bank aussah. Dort gab sie ihre Kennung, ihr Passwort und einen per SMS empfangenen Code ein. Sie erklärt ausdrücklich, keine Überweisung beauftragt, keinen Empfänger ausgewählt und keine Zahlung über 700,00 EUR gebilligt zu haben. Welche technische Authentifizierung der spätere Zahlungsauftrag durchlief, ist unbekannt. Die Kanzlei darf aus der Eingabe der Zugangsdaten keine zusätzliche Freigabehandlung ergänzen.

Am Dienstag, 13.10.2026, 08.42 Uhr, meldete sie die Belastung der Bank. Am 14.10. lehnte die Bank eine Erstattung mit der Weitergabe der Zugangsdaten und grober Fahrlässigkeit ab. Die Bank benennt keinen eigenen Betrugsverdacht gegen die Kundin. Bis zum Auszug vom 16.10., 08.00 Uhr, erfolgte keine Gutschrift. Die Mitarbeiterin im Rückrufdienst erklärte, der Betrag könne nach Abschluss einer internen Prüfung zurückkommen; eine Zusage liegt darin nicht.

3. Erwartetes Ergebnis

Die Kanzlei soll die sofort verlangte Erstattung, die technische Aufklärung und ein mögliches späteres Ersatzverlangen der Bank getrennt bearbeiten. Ein Entwurf des Erstattungsschreibens und ein noch nicht freigegebener Klageentwurf liegen bei. Die Anlagen K1 bis K8 sollen belegt zugeordnet werden. Frau Fenchel entscheidet nach Rücksprache mit Frau Rabenstein über Versand oder Einreichung. Offene Angaben sollen als Fragen sichtbar bleiben; fehlende Bankprotokolle dürfen nicht erfunden werden.''')
    add(2,'docx','Gedaechtnisprotokoll_Rabenstein', '''Kunigunde Rabenstein. Niederschrift vom 15.10.2026, ergänzt nach telefonischer Rückfrage am 16.10.2026, 09.35 Uhr.

1. Der angebotene Sessel

Ich hatte einen grünen Lesesessel eingestellt, weil meine Schwester umzieht und mir einen kleineren mitbringt. Der Preis war 85 Euro, Abholung in Halle. Der Interessent nannte sich „Lars B.“ und schrieb, er würde eine Abholung organisieren. Das kam mir etwas umständlich vor. Ich habe aber schon öfter gebrauchte Sachen abgegeben und wollte die Unterhaltung neben dem Abendessen zu Ende bringen. Der Sessel steht noch bei mir. Niemand hat ihn abgeholt, und ich habe auch keine 85 Euro bekommen.

2. Die vermeintliche Bestätigung

Es kam ein Hinweis, ich müsse den Verkauf bestätigen. Ich weiß noch, dass ich eine Seite sah, die farblich und mit dem Schriftzug nach meiner Bank aussah. Ich gab meine Kennung und mein Passwort sowie den danach per SMS empfangenen Code ein. Den vollständigen Wortlaut der SMS habe ich nicht gesichert. Danach bekam ich keine vernünftige Bestätigung. Ich habe nichts über einen Empfänger namens N. Kramer gelesen und keine 700 Euro eingetragen. Ich wollte Geld erhalten, nicht Geld überweisen.

Außer der Eingabe dieses SMS-Codes kann ich zu weiteren Bildschirmen oder einer technischen Sicherheitsabfrage heute nichts Verlässliches sagen. Insbesondere erinnere ich mich nicht sicher an den vollständigen Nachrichtentext, einen Empfänger oder einen ausgewiesenen Betrag. Ich möchte nicht etwas bestätigen, nur weil die Bank fragt, ob ich eine Freigabe vorgenommen habe. Ich habe jedenfalls keinen Zahlungsauftrag erteilen wollen. Mein Telefon ist noch vorhanden. Ich habe seit dem Vorfall nur mit dem Kundendienst gesprochen und die von der Bank veranlasste Zugangssperre abgewartet. Zugangsdaten habe ich nicht an die Kanzlei geschickt.

3. Entdeckung und Meldung

Am Dienstagmorgen habe ich nachgesehen, ob die 85 Euro da sind. Stattdessen sah ich die 700-Euro-Belastung. Ich rief die Bank an. Meine Handy-Anrufliste zeigt den Beginn um 08.42 Uhr; eine eigene Aufzeichnung des Gesprächs habe ich nicht. Der Mitarbeiter sagte mir eine Sperre und eine Vorgangsnummer zu. Den Auszug und den Chat habe ich danach gespeichert. Ich hatte meiner Schwester zunächst gesagt, es sei „am Wochenende“ passiert. Nach dem gespeicherten Chat und meinem Dienstplan war es Montag, der 12. Oktober. Die Uhrzeit 19.17 Uhr stammt aus dem Bankauszug, nicht aus meiner Erinnerung.

4. Was ich erreichen möchte

Ich kann die laufenden Ausgaben derzeit noch bezahlen. Trotzdem fehlen mir 700 Euro, und ich verstehe nicht, weshalb die Bank das so einfach ablehnt. Ich möchte keine riskante Auseinandersetzung über etwas, das ich angeblich eingeräumt haben soll. Bitte nehmen Sie nur das in ein Schreiben auf, was ich tatsächlich gesagt habe. Ich bin bereit, mein Telefon fachgerecht untersuchen zu lassen, wenn vorher geklärt ist, welche Daten benötigt werden und was das kostet.

Kunigunde Rabenstein. Der Text ist von der Mandantin per E-Mail bestätigt; eine handschriftliche Unterzeichnung ist noch nicht erfolgt.''')
    add(3,'docx','Telefonnotiz_Bankkontakt', '''Kanzlei Fenchel. Telefonvermerk vom 16.10.2026, 10.05 bis 10.17 Uhr. Gesprächspartner: Daria Fenchel und Ottmar Wenzel, Reklamationsbearbeitung der Saaleufer Bank AG. Bankvorgang SB-2026-1013-4821.

Frau Fenchel nennt die Mandatsanzeige und fragt nach dem Stand der Erstattung. Herr Wenzel bestätigt, dass die Meldung am 13.10.2026 um 08.42 Uhr erfasst wurde und der Zugang seit 08.51 Uhr gesperrt ist. Die Wertstellung der Belastung ist nach seiner Bildschirmansicht der 12.10.2026. Eine Erstattung ist noch nicht gebucht. Die Aussage betrifft nur den Bearbeitungsstand, nicht das Ergebnis einer technischen Prüfung.

Auf die Frage nach dem Autorisierungsnachweis erklärt Herr Wenzel, dass der Vorgang als Überweisung im System stehe. Er könne aus seinem Arbeitsbereich keine vollständigen Anmelde-, Geräte-, Authentifizierungs- oder Freigabeprotokolle ausgeben. Frau Fenchel weist darauf hin, dass eine Buchungsart allein die Erklärung der Kundin nicht beantwortet. Herr Wenzel sagt zu, die Anfrage an die zuständige Stelle weiterzugeben. Er nennt keinen feststehenden Termin für deren Antwort.

Frau Fenchel fragt, ob die Bank Frau Rabenstein selbst betrügerisches Handeln vorwirft oder einen solchen Verdacht an eine zuständige Behörde gemeldet hat. Herr Wenzel erklärt: „Bei uns steht grob fahrlässig, weil die Daten auf der fremden Seite eingegeben wurden. Einen Betrugsvorwurf gegen Ihre Mandantin habe ich hier nicht.“ Eine verbindliche Erklärung der Rechtsabteilung über sämtliche Meldungen kann er nicht abgeben. Eine schriftliche Antwort wird angefordert.

Herr Wenzel möchte die Sache zunächst bis Montag, 19.10.2026, intern prüfen lassen. Frau Fenchel sagt keine Stillhaltevereinbarung zu. Sie kündigt lediglich einen zu prüfenden schriftlichen Entwurf an. Weder ein Anerkenntnis grober Fahrlässigkeit noch ein Vergleich noch ein Verzicht auf sofortige Erstattung wird erklärt. Gesprächsnotiz erstellt um 10.22 Uhr; die Bank hat den Vermerk nicht bestätigt.''')
    add(4,'docx','Erstattungsschreiben_Entwurf', '''Rechtsanwältin Daria Fenchel, Fiktiver Birkenbogen 8, 06108 Halle (Saale). An die Saaleufer Bank AG, Reklamationsbearbeitung, Fiktiver Uferplatz 4, 06108 Halle (Saale). 16.10.2026. Unser Zeichen DF 166/26; Ihr Vorgang SB-2026-1013-4821. Entwurf, noch nicht versandt.

Erstattung der nicht autorisierten Belastung über 700 Euro

Sehr geehrte Damen und Herren,

ich vertrete Frau Kunigunde Rabenstein. Meine Mandantin verlangt die Erstattung der am 12.10.2026 gebuchten, von ihr nicht autorisierten Zahlung über 700,00 EUR und die entsprechende Wiederherstellung ihres Zahlungskontos KR-4821. Sie zeigte Ihnen den Vorgang am 13.10.2026 um 08.42 Uhr an. Ihre Eingangsbestätigung und Ihr Schreiben vom 14.10.2026 liegen bei. Aus dem Auszug vom heutigen Morgen ergibt sich keine Erstattung. Belege: Anlagen K1, K4, K5 und K6.

Meine Mandantin wollte über Hofmarkt einen Sessel für 85,00 EUR verkaufen. Ein vermeintlicher Käufer führte sie zu einer bankähnlichen Seite. Sie gab dort Zugangsdaten und einen per SMS erhaltenen Code ein. Sie erteilte dadurch nach ihrer Darstellung keinen Zahlungsauftrag und billigte die Überweisung an den im Auszug genannten Empfänger nicht. Den Verkaufserlös hat sie nicht erhalten. Belege: Anlagen K2, K3 und K7. Ob und wie Ihr System den späteren Auftrag technisch authentifizierte, ist aufzuklären; eine zusätzliche Zahlungsfreigabe meiner Mandantin ist nicht belegt.

§ 675u Sätze 2 und 3 BGB verlangt Erstattung und Kontowiederherstellung unverzüglich, spätestens bis zum Ende des folgenden Geschäftstags. Nach Ihrer Bestätigung der Anzeige am 13.10.2026 und Ihrer Geschäftstagsauskunft war dies hier spätestens Mittwoch, der 14.10.2026. Ihr Ablehnungsschreiben beruft sich allein auf grobe Fahrlässigkeit. Der Gerichtshof der Europäischen Union hat am 15.10.2026 in der Rechtssache C-70/25, Tukowiecka, zwischen der unverzüglichen Erstattung eines nicht autorisierten Zahlungsvorgangs und der späteren Verteilung des Verlusts unterschieden. Nach Rn. 31 bis 34 und 39 bis 42 dieser Entscheidung gestattet der bloße Vorwurf grober Fahrlässigkeit keine Verweigerung der unverzüglichen Erstattung. Die besondere Ausnahme setzt einen begründeten Verdacht eigenen betrügerischen Handelns des Zahlers und die schriftliche Mitteilung der Gründe an die zuständige Behörde voraus. Die Entscheidung betrifft die Auslegung von Artikel 73 Absatz 1 und Artikel 74 Absatz 1 der Richtlinie (EU) 2015/2366. Ein späteres, begründetes Ersatzverlangen wäre gesondert zu prüfen; meine Mandantin erkennt ein solches nicht an.

Bitte stellen Sie das Konto unverzüglich entsprechend dem Erstattungsanspruch wieder her und bestätigen Sie den Vollzug schriftlich. Diese Aufforderung enthält keine Verlängerung einer bereits laufenden oder abgelaufenen gesetzlichen Erstattungsfrist. Soweit Sie sich auf einen rechtlich erheblichen Betrugsverdacht gegen meine Mandantin berufen wollen, teilen Sie bitte die konkreten Gründe und die dazu erfolgte behördliche Unterrichtung mit. Die bloße Weitergabe von Zugangsdaten ist in Ihrem Schreiben bislang nur als Fahrlässigkeitsvorwurf bezeichnet.

Bitte sichern Sie außerdem die fallbezogenen Anmelde-, Geräte-, Authentifizierungs- und Zahlungsdaten gegen routinemäßige Löschung und teilen Sie mit, welche aussagekräftigen Nachweise Sie zur Verfügung stellen können. Diese Bitte behauptet keinen grenzenlosen Anspruch auf interne Systeme oder Daten unbeteiligter Personen. Den Zugangssperrvermerk haben wir als Anlage K4 erfasst; den vollständigen Zahlungsnachweis benötigen wir weiterhin.

Mit freundlichen Grüßen
Daria Fenchel, Rechtsanwältin

Anlagen: K1 Belastungsbeleg; K2 Inserat; K3 Nachrichtenverlauf; K4 Eingangsbestätigung und Sperrvermerk; K5 Ablehnung; K6 Folgeauszug; K7 bestätigtes Gedächtnisprotokoll; K8 Vollmacht. Die Anlagenfolge ist vor Versand mit der Arbeitsmappe abzugleichen.''')
    add(5,'docx','Klageschrift_Entwurf', '''Rechtsanwältin Daria Fenchel, Fiktiver Birkenbogen 8, 06108 Halle (Saale). An das Amtsgericht Halle (Saale). Entwurf vom 16.10.2026, noch nicht eingereicht. Akte DF 166/26.

Klage wegen einer nicht autorisierten Zahlung

Kunigunde Rabenstein, Fiktive Lindenhofgasse 17, 06108 Halle (Saale), Klägerin, Prozessbevollmächtigte Rechtsanwältin Daria Fenchel, gegen Saaleufer Bank AG, Fiktiver Uferplatz 4, 06108 Halle (Saale), Beklagte, gesetzlich vertreten durch [vertretungsberechtigten Vorstand nach verifiziertem Registerstand ergänzen].

1. Anträge

Die Beklagte wird verurteilt, der Klägerin den am 12.10.2026 als Zahlung an N. Kramer abgebuchten Betrag von 700,00 EUR zu erstatten und das bei ihr geführte Zahlungskonto mit der Kundenkennung KR-4821 insoweit auf den Stand zu bringen, auf dem es sich ohne diese Belastung befunden hätte. [Die Kontobezeichnung und die prozessual eindeutig gefasste Wiederherstellung vor Einreichung mit dem aktuellen Kontoauszug abgleichen.] Die Beklagte trägt die Kosten des Rechtsstreits.

Ein Zinsantrag und weitere Vermögensnachteile werden in diesem Arbeitsstand nicht beziffert. Vor Einreichung ist anhand des Erstattungsverlaufs und der konkreten Fälligkeit zu entscheiden, ob und in welcher Form sie geltend gemacht werden. Die Klage ist nicht bereits deshalb freigegeben, weil die Hauptsumme feststeht.

2. Konto und streitige Belastung

Die Klägerin führt bei der Beklagten ein privates Zahlungskonto. Am 12.10.2026 wurde dieses mit 700,00 EUR zugunsten des Empfängers N. Kramer belastet. Der Bankauszug nennt die Buchungskennung ZV-121026-1917 und die Wertstellung 12.10.2026. Die Klägerin bestreitet, diesen Zahlungsvorgang autorisiert zu haben. Beweis: Belastungsbeleg, Anlage K1; [vollständigen Zahlungsdiensterahmenvertrag und aktuellen Kontonachweis nach Eingang genau bezeichnen].

Die Klägerin bot auf Hofmarkt einen gebrauchten Sessel für 85,00 EUR an. Ein Nutzer unter dem Anzeigenamen „Lars B.“ gab sich als Kaufinteressent aus. Über einen vermeintlichen Bestätigungslink gelangte die Klägerin auf eine bankähnliche Seite und gab dort ihre Zugangsdaten und einen per SMS erhaltenen Code ein. Eine Überweisung über 700,00 EUR wollte und billigte sie nicht. Die belastete Zahlung wurde durch einen Dritten veranlasst. Beweis: Inserat, Anlage K2; gesicherter Nachrichtenverlauf, Anlage K3; ergänzende persönliche Anhörung der Klägerin. Das eigene Gedächtnisprotokoll, Anlage K7, dokumentiert ihre Darstellung und wird nicht als unabhängiger technischer Beweis ausgegeben.

3. Anzeige und Ablehnung

Am 13.10.2026, 08.42 Uhr, zeigte die Klägerin die Zahlung der Beklagten an. Die Beklagte bestätigte den Eingang und die Zugangssperre. Am 14.10.2026 lehnte sie die Erstattung ab, weil die Klägerin Zugangsdaten auf einer fremden Seite eingegeben und dadurch grob fahrlässig gehandelt habe. Ein eigener Betrugsverdacht gegen die Klägerin wurde nicht benannt. Eine Wiederherstellung erfolgte bis zum Auszug vom 16.10.2026, 08.00 Uhr, nicht. Beweis: Anlagen K4 bis K6. [Den Stand unmittelbar vor Einreichung durch einen weiteren Auszug aktualisieren; eine spätere Erstattung nicht übergehen.]

4. Rechtliche Begründung

Die Klägerin stützt ihr Begehren auf § 675u Sätze 2 und 3 BGB. Bei Anzeige am 13.10.2026 und dem von der Bank bestätigten Geschäftstag 14.10.2026 war die Wiederherstellung spätestens bis Ende des 14.10.2026 vorzunehmen. Die Vorschrift ist im Zusammenhang mit Artikel 73 der Richtlinie (EU) 2015/2366 auszulegen. Das Urteil des Gerichtshofs der Europäischen Union vom 15.10.2026, C-70/25, Tukowiecka, betrifft gerade die Weigerung eines Zahlungsdienstleisters, wegen behaupteter grober Fahrlässigkeit sofort zu erstatten. Die unverzügliche Erstattung ist nach Rn. 31 bis 34 spätestens bis zum Ende des folgenden Geschäftstags vorzunehmen; der Vorwurf grober Fahrlässigkeit eröffnet keine eigene Ausnahme. Die Ausnahme des begründeten Betrugsverdachts betrifft eigenes betrügerisches Handeln des Zahlers und verlangt die schriftliche Mitteilung der Gründe an die zuständige Behörde, Rn. 33 und 39. Das Urteil trennt diese erste Erstattungsstufe von der späteren Verlustverteilung, Rn. 38 und 42. Eine abschließende Haftungsfreiheit der Klägerin wird damit nicht behauptet, Rn. 46 und 52; EuGH, Urteil vom 15.10.2026, C-70/25, Tukowiecka, ECLI:EU:C:2026:852. Die behauptete grobe Fahrlässigkeit darf nicht ohne Weiteres an die Stelle der Voraussetzungen einer Ausnahme von der unverzüglichen Erstattung treten.

Damit ist die endgültige Verlustverteilung nicht vorweggenommen. Sollte die Beklagte ein späteres Ersatzverlangen erheben, sind dessen Voraussetzungen, die konkret verletzte Pflicht, der Verschuldensmaßstab und der Zusammenhang mit dem Verlust eigenständig zu prüfen. Die Klägerin bestreitet ein pauschal aus der Zugangsdatenverwendung abgeleitetes Anerkenntnis. Die technische Erfassung eines Auftrags ersetzt nicht die Prüfung, ob die Klägerin die streitige Zahlung autorisierte. Nach § 675w Satz 1 BGB hat die Beklagte die Authentifizierung sowie die ordnungsgemäße, störungsfreie Erfassung und Verbuchung nachzuweisen. Die bloße Nutzungsaufzeichnung beweist nach Satz 3 nicht notwendigerweise die Autorisierung. Satz 4 verlangt unterstützende Beweismittel für einen Betrugs-, Vorsatz- oder groben Fahrlässigkeitsvorwurf. Für ein späteres Ersatzverlangen nach § 675v Absatz 3 BGB sind der kausale Pflichtverstoß und das Verschulden gesondert darzulegen; Absatz 4 ist nach Aufklärung der starken Kundenauthentifizierung mitzuprüfen. [Etwaige konkret erklärte Gegenrechte anhand des späteren Bankvortrags einarbeiten.]

5. Noch benötigte Belege und Anlagen

Die vollständigen Bankprotokolle zu Anmeldung, Gerät, Authentifizierung und Freigabe liegen nicht vor. Die Klägerin berichtet von der Eingabe eines per SMS erhaltenen Codes auf der nachgeahmten Seite. Dessen vollständiger Begleittext und technische Funktion sind nicht gesichert; daraus wird keine Autorisierung der streitigen Zahlung abgeleitet. Nach Eingang der Stellungnahme der Beklagten ist zu konkretisieren, welche entscheidungserheblichen Tatsachen bestritten werden und welche prozessual zulässige Beweisaufnahme hierfür beantragt werden soll. [Ein konkreter Antrag auf Urkundenvorlage setzt die zu bezeichnende Unterlage und die Verfahrensvoraussetzungen voraus; vor Einreichung ausformulieren.]

Anlagen K1 bis K8 entsprechen der Zuordnung im Erstattungsentwurf. K8 ist die Vollmacht; ihre Vorlage beantwortet keine Frage zur Zahlungsautorisierung. Gerichtliche Zuständigkeit, Vertretung der Beklagten, vollständiger Antrag, aktueller Kontostand, Kosten und anwaltliche Freigabe sind vor Einreichung zu prüfen.

Daria Fenchel, Rechtsanwältin''')
    add(6,'docx','Vollmacht_Rabenstein', '''Kunigunde Rabenstein, Fiktive Lindenhofgasse 17, 06108 Halle (Saale). Erklärung vom 16.10.2026.

Ich bevollmächtige Rechtsanwältin Daria Fenchel, mich wegen der am 12.10.2026 erfolgten Belastung meines privaten Zahlungskontos bei der Saaleufer Bank AG über 700,00 EUR zu vertreten. Sie darf die hierfür erforderlichen Unterlagen anfordern, Erklärungen entgegennehmen und nach Abstimmung die Erstattung geltend machen. Ein Vergleich, ein Verzicht oder ein Anerkenntnis einer Ersatzforderung der Bank bedarf meiner vorherigen Zustimmung.

Die Vorbereitung eines Klageentwurfs ist beauftragt. Über die Einreichung möchte ich nach Erläuterung der Kosten und Erfolgsaussichten gesondert entscheiden. Diese Vollmacht enthält keine Erlaubnis, mit meinen Zugangsdaten auf das Bankkonto zuzugreifen, Zahlungen auszulösen oder ein Sicherheitsverfahren zu bedienen. Meine Passwörter und sonstigen Geheimnisse sollen nicht in der Kanzleiakte gespeichert werden.

Organisatorische Nachrichten dürfen an kunigunde.rabenstein@postfach.example gesandt werden. Meine Schwester Ottilie darf über Termine informiert werden, erhält aber ohne weitere Absprache keine Kontoauszüge. Ich werde Änderungen meiner Kontaktdaten mitteilen und etwaige Erstattungen sofort an die Kanzlei melden.

Kunigunde Rabenstein. Die Erklärung ist im Mandatsgespräch bestätigt worden. Eine handschriftliche Unterschrift ist in dieser Datei nicht abgebildet.''')
    add(7,'pdf','Belastungsbeleg_12_Oktober', '''Saaleufer Bank AG. Elektronischer Umsatzbeleg, abgerufen am 13.10.2026, 08.36 Uhr. Kontoinhaberin Kunigunde Rabenstein; Kundenkennung KR-4821. Empfängeranschrift der Kontoinhaberin: Fiktive Lindenhofgasse 17, 06108 Halle (Saale).

1. Buchung

Buchungstag: 12.10.2026. Wertstellung: 12.10.2026. Betrag: minus 700,00 EUR. Buchungstext: Überweisung N. Kramer. Interne Kennung: ZV-121026-1917. Im Umsatzabruf angezeigte Erfassungszeit: 19.17 Uhr. Empfängerkonto im Kundenauszug nur gekürzt angezeigt: Ende 9054. Es ist keine vollständige verwendbare IBAN abgedruckt.

Kontostand unmittelbar vor der Belastung: 2.146,38 EUR. Kontostand unmittelbar nach der Belastung: 1.446,38 EUR. Die Anzeige umfasst diesen einzelnen Umsatz. Sie enthält weder ein vollständiges Sitzungsprotokoll noch Angaben zum Endgerät, zum Sicherheitsverfahren oder zu einer Freigabehandlung. Die Uhrzeit im Buchungssystem ist nicht der Nachweis einer Erklärung der Kundin.

2. Abrufhinweis

Dieser Umsatzbeleg wurde von Frau Rabenstein nach Entdeckung der Belastung gespeichert. Eine Rückerstattung oder Rückbelastung ist auf dieser Seite nicht ausgewiesen. Die Empfängerbezeichnung ist die des Umsatzdatensatzes; die Identität der dahinterstehenden Person wurde von der Kundin nicht überprüft. Rückfragen zum Buchungssatz nimmt die Reklamationsbearbeitung unter Angabe der internen Kennung entgegen.''')
    add(8,'pdf','Inserat_Hofmarkt', '''Hofmarkt. Gespeicherte Ansicht des eigenen Inserats von Kunigunde Rabenstein. Abruf am 13.10.2026, 09.18 Uhr. Anzeigenkennung HM-HAL-1210-85.

Grüner Lesesessel abzugeben

Preis: 85,00 EUR. Standort: Halle (Saale), Abholung nach Vereinbarung. Eingestellt am Montag, 12.10.2026, 17.58 Uhr. Zustand: gebraucht, stabil, am linken Arm etwas ausgeblichen. Breite ungefähr 74 cm. Der Sessel muss eine Treppe heruntergetragen werden; bei der Terminabsprache kann geklärt werden, ob jemand helfen kann.

Anbieterin: Kunigunde R. Die öffentliche Anzeige enthält keine Bankdaten und keine vollständige Wohnanschrift. Im Verkäuferkonto ist der Status am 13.10. weiterhin „online“. Die Plattform weist keine über sie abgeschlossene Zahlung und keinen bestätigten Versandauftrag aus. Das ist eine Anzeige aus dem Verkäuferkonto, keine Auskunft über sämtliche Vorgänge außerhalb der Plattform.

Notiz der Anbieterin beim Speichern: Der Sessel ist noch da. Ich habe den angeblichen Käufer nach dem Vorfall nicht erneut kontaktiert. Ich möchte das Inserat für die Aufklärung sichern und anschließend entfernen lassen.''')
    add(9,'pdf','Nachrichtenverlauf_Hofmarkt', '''Hofmarkt. Vom Verkäuferkonto gesicherter Nachrichtenverlauf, Unterhaltung zu HM-HAL-1210-85. Export am 13.10.2026 um 09.24 Uhr. Benutzerangaben wurden von der Plattform nicht als echte Identitäten bestätigt.

12.10.2026, 18.34 Uhr, „Lars B.“: Guten Abend, ist der grüne Sessel noch da? Ich könnte jemanden zum Abholen schicken.

12.10.2026, 18.41 Uhr, Kunigunde R.: Ja, noch da. Am Mittwoch wäre ich ab 17 Uhr zu Hause. Bitte mit zwei Leuten, ich kann nicht schwer tragen.

12.10.2026, 18.48 Uhr, „Lars B.“: Passt. Ich kümmere mich um die Abholung. Sie erhalten eine Bestätigung, damit alles zugeordnet ist.

12.10.2026, 19.03 Uhr, Kunigunde R.: Ich sehe jetzt einen Hinweis zur Bestätigung. Ist das von Hofmarkt?

12.10.2026, 19.05 Uhr, „Lars B.“: Ja, das gehört dazu.

12.10.2026, 19.22 Uhr, Kunigunde R.: Es ist nichts bestätigt worden. Ich habe auch keinen Termin erhalten. Können Sie mir bitte schreiben, wer am Mittwoch kommt?

13.10.2026, 08.31 Uhr, Kunigunde R.: Auf meinem Konto fehlen 700 Euro. Was ist hier passiert?

Im gesicherten Verlauf ist danach keine Antwort enthalten. Der gefährliche Link selbst wurde nicht in diese Arbeitskopie übernommen. Eine einzelne externe Zieladresse ist hier nicht enthalten. Der Verlauf zeigt Nachrichten und Zeitstempel; er belegt keine technische Zahlungsfreigabe. Der Name „Lars B.“ und der Zahlungsempfänger N. Kramer sind verschiedene Bezeichnungen, deren Zusammenhang noch nicht aufgeklärt ist.''')
    add(10,'pdf','Eingang_Bank_und_Sperrung', '''Saaleufer Bank AG. Vorgangsbestätigung vom 13.10.2026, 09.02 Uhr. Vorgang SB-2026-1013-4821. An Frau Kunigunde Rabenstein.

Sehr geehrte Frau Rabenstein,

wir bestätigen Ihre telefonische Meldung vom heutigen Tag, 08.42 Uhr, zur Belastung über 700,00 EUR mit der Kennung ZV-121026-1917. Sie haben erklärt, diese Zahlung nicht beauftragt zu haben. Der Vorgang wurde zur Prüfung aufgenommen. Ihren Onlinezugang haben wir um 08.51 Uhr gesperrt. Diese Sperre wirkt auf den Zugang und sagt nichts über den Ausgang der Reklamation aus.

Bitte verwenden Sie den verdächtigen Link nicht erneut und übersenden Sie uns keine Passwörter. Den Umsatzbeleg und vorhandene Nachrichten können Sie dem Vorgang zuordnen lassen. Ob ein Rückruf der Zahlung möglich ist, wird gesondert geprüft. Eine Rückholung ist mit diesem Schreiben nicht bestätigt.

Nach unseren Kontobedingungen sind Montag bis Freitag, soweit am Sitz unserer Bank kein gesetzlicher Feiertag ist, Geschäftstage für diese Zahlungsdienstleistungen. Am 13. und 14.10.2026 war unsere Reklamationsbearbeitung regulär besetzt. Den Inhalt Ihrer Anzeige haben wir im Vorgang gespeichert. Eine abschließende technische Bewertung ist darin noch nicht enthalten.

Mit freundlichen Grüßen
Saaleufer Bank AG, Reklamationsbearbeitung''')
    add(11,'pdf','Ablehnung_Saaleufer_Bank', '''Saaleufer Bank AG, Fiktiver Uferplatz 4, 06108 Halle (Saale). An Frau Kunigunde Rabenstein, Fiktive Lindenhofgasse 17, 06108 Halle (Saale). 14.10.2026. Vorgang SB-2026-1013-4821.

Ihre Reklamation über 700 Euro

Sehr geehrte Frau Rabenstein,

wir haben Ihre Schilderung ausgewertet. Sie haben Ihre persönlichen Zugangsdaten auf einer fremden Internetseite eingegeben, obwohl diese Daten geheim zu halten sind. Wir bewerten dies als grob fahrlässige Verletzung Ihrer Sorgfaltspflichten. Durch die Preisgabe konnten Dritte die Zahlung auslösen. Aus diesem Grund erstatten wir den Betrag von 700,00 EUR nicht.

Unsere technische Fachstelle ist noch mit der Zusammenstellung der verfügbaren Vorgangsdaten befasst. Wir übermitteln Ihnen mit diesem Schreiben kein vollständiges Anmelde- oder Freigabeprotokoll. Der Kontoumsatz wird in unserem System als Überweisung geführt. Falls Sie weitere Unterlagen einreichen möchten, geben Sie bitte unsere Vorgangsnummer an.

Die Zugangssperre bleibt bestehen, bis die Wiederherstellung mit Ihnen abgestimmt ist. Die Prüfung einer möglichen Rückholung vom Empfängerkonto dauert an. Damit ist keine Zusage einer Gutschrift verbunden.

Mit freundlichen Grüßen
Ottmar Wenzel, Reklamationsbearbeitung''')
    add(12,'pdf','Folgeauszug_bis_16_Oktober', '''Saaleufer Bank AG. Kontoumsätze für Kunigunde Rabenstein, Kundenkennung KR-4821. Zeitraum 12.10.2026 bis 16.10.2026, 08.00 Uhr. Abruf im nach Identitätsprüfung bereitgestellten Dokumentenbereich.

Stand vor der streitigen Buchung am 12.10.: 2.146,38 EUR.

12.10.2026, Überweisung N. Kramer, ZV-121026-1917: minus 700,00 EUR; Wertstellung 12.10.2026; Saldo 1.446,38 EUR.

13.10.2026, Lastschrift Stadtwerke, Abschlag Oktober: minus 62,00 EUR; Wertstellung 13.10.2026; Saldo 1.384,38 EUR.

14.10.2026, Kartenzahlung Lebensmittel: minus 27,45 EUR; Wertstellung 14.10.2026; Saldo 1.356,93 EUR.

15.10.2026, Erstattung Apotheke: plus 12,80 EUR; Wertstellung 15.10.2026; Saldo 1.369,73 EUR.

16.10.2026, 08.00 Uhr: keine weitere Buchung. Endsaldo 1.369,73 EUR.

Die Apothekenerstattung betrifft einen Kassenbon vom 09.10.2026 und hat keinen Bezug zur streitigen Zahlung. Im ausgewiesenen Zeitraum wurde der Betrag von 700,00 EUR nicht erstattet. Es sind keine gesonderten Rückrufkosten oder Überziehungszinsen gebucht. Die Umsatzliste kann nur den genannten Zeitraum abbilden; nachfolgende Buchungen sind vor einem endgültigen Anspruchsschreiben erneut abzugleichen.''')
    add(13,'pdf','Plattformantwort_15_Oktober', '''Hofmarkt Kundenservice. Nachricht zum Vorgang HM-SICH-4821 vom 15.10.2026, 11.20 Uhr. An Kunigunde Rabenstein.

Sehr geehrte Frau Rabenstein,

wir haben Ihre Meldung zum Nutzerkonto „Lars B.“ erhalten und den übermittelten Nachrichtenverlauf dem Sicherheitsvorgang zugeordnet. Das Inserat betrifft einen Lesesessel für 85,00 EUR. In den uns für Ihre Verkäuferansicht vorliegenden Daten ist keine bestätigte Zahlung für dieses Inserat ausgewiesen. Bitte unterscheiden Sie zwischen Nachrichten eines Nutzers und Mitteilungen unseres Kundenservice.

Für die Bestätigung einer Abholung brauchen Sie unserem Kundenservice kein Onlinebanking-Passwort mitzuteilen. Wir haben den von Ihnen beschriebenen externen Vorgang nicht selbst beobachtet. Eine Aussage darüber, wie der Zahlungsauftrag bei Ihrer Bank technisch ausgelöst wurde, können wir daher nicht treffen.

Der von Ihnen gemeldete Nutzerzugang ist derzeit eingeschränkt. Das bedeutet weder, dass wir seine bürgerliche Identität bestätigt haben, noch dass wir den Namen aus Ihrem Bankumsatz zuordnen können. Ihre Nachricht und die gespeicherten Inhalte werden im Sicherheitsvorgang vorgehalten. Für weitergehende Datenübermittlungen benötigen wir den konkreten Zweck und den zulässigen Verfahrensweg.

Mit freundlichen Grüßen
Nele Dorn, Hofmarkt Kundenservice''')
    add(14,'pdf','Strafanzeige_Sachverhalt_13_Oktober', '''Kunigunde Rabenstein, Fiktive Lindenhofgasse 17, 06108 Halle (Saale). Private Kopie der Sachverhaltsschilderung vom 13.10.2026. Nach eigener Erklärung am selben Tag um 12.10 Uhr an die Polizei übermittelt. Ein polizeiliches Aktenzeichen und eine amtliche Eingangsbestätigung liegen der Kanzlei noch nicht vor.

Ich möchte anzeigen, dass am 12.10.2026 eine fremde Person 700,00 EUR von meinem Konto bei der Saaleufer Bank AG überwiesen hat. Vorher hatte sich auf Hofmarkt ein vermeintlicher Käufer meines Sessels gemeldet. Über einen angeblichen Bestätigungsvorgang gelangte ich auf eine Seite, die wie meine Bank aussah. Ich gab dort meine Zugangsdaten und einen per SMS erhaltenen Code ein. Ich wollte keine Zahlung leisten und habe keine Überweisung über 700,00 EUR gebilligt.

Der Interessent nannte sich „Lars B.“. Im Kontoumsatz steht dagegen N. Kramer. Ob diese Bezeichnungen dieselbe Person meinen, weiß ich nicht. Ich kenne weder den tatsächlichen Namen noch eine gesicherte Anschrift des Handelnden. Die 700,00 EUR sind nicht der Preis des Sessels. Dieser war für 85,00 EUR angeboten und wurde nicht abgeholt.

Am 13.10.2026 um 08.42 Uhr habe ich den Vorfall meiner Bank gemeldet. Die Bank sperrte den Zugang und bestätigte die Meldung. Eine Erstattung erfolgte bis zum letzten Auszug nicht. Mein Telefon sowie der gespeicherte Nachrichtenverlauf sind vorhanden. Ich habe keine eigene Untersuchung des Geräts vorgenommen und möchte den ursprünglichen Zustand der vorhandenen Belege möglichst erhalten.

Beifügen möchte ich den Belastungsbeleg, das Inserat, den Nachrichtenverlauf sowie die Eingangsbestätigung der Bank. Passwörter sind darin nicht enthalten. Bitte teilen Sie mir nach Eingang mit, unter welchem Aktenzeichen die Sache geführt wird und welche weiteren Unterlagen benötigt werden. Diese Niederschrift enthält meine Erinnerung; die technischen Abläufe kann ich nicht selbst erklären.

Kunigunde Rabenstein''')
    add(15,'pdf','Geraet_und_Belegsicherung', '''Kanzlei Fenchel. Übernahmevermerk vom 16.10.2026, 09.50 Uhr. Gegenstand: Telefon und vorhandene Dateien von Kunigunde Rabenstein.

1. Vorhandene Unterlagen

Die Mandantin zeigt ihr eigenes, weiterhin benutztes Telefon. Eine forensische Kopie wird heute nicht erstellt. Übergeben werden der gespeicherte Umsatzbeleg, das Inserat und der Nachrichtenverlauf in der vorhandenen Exportansicht. Das Datum der Exporte ist in den Dateien vermerkt. Die Kanzlei bestätigt mit diesem Vermerk nur den Eingang, nicht die Unveränderlichkeit aller ursprünglichen Plattformdaten.

2. Offene technische Angaben

Es liegt kein vollständiger Browser- oder Sicherheitsverlauf vor. Die tatsächliche Zieladresse der nachgeahmten Seite wird nicht erneut geöffnet. Es werden keine Zugangsdaten notiert. Die Mandantin berichtet die Eingabe eines SMS-Codes, kann dessen vollständigen Begleittext und weitere Sicherheitsbildschirme aber nicht zuverlässig rekonstruieren. Über TAN-Verfahren, Gerätekopplung, starke Kundenauthentifizierung, Empfängeranzeige oder einen behaupteten Freigabeschritt darf aus diesem Vermerk keine Feststellung abgeleitet werden.

3. Weiteres Vorgehen

Die vorhandenen Dateien bleiben unverändert erhalten. Neue Auskünfte sollen daneben mit Datum und Herkunft abgelegt werden. Für eine Untersuchung des Telefons sind zunächst Umfang, Datenschutz und Kosten abzustimmen. Die Bank soll ihre fallbezogenen Protokolle sichern. Eine Anfrage nach Belegen ist noch kein Nachweis, dass diese Belege existieren oder bereits herausgegeben wurden.

Daria Fenchel''')
    add(16,'pdf','Rueckrufstatus_und_Bankbedingungen', '''Saaleufer Bank AG. Ergänzende Vorgangsmitteilung vom 16.10.2026, 11.30 Uhr. Vorgang SB-2026-1013-4821.

Sehr geehrte Frau Rabenstein,

die Anfrage zur Rückholung der Zahlung vom 12.10.2026 ist noch offen. Eine Zustimmung des Empfängers oder eine erfolgreiche Rückbuchung wurde uns bisher nicht bestätigt. Wir haben Ihnen keine Rückrufkosten berechnet. Sollte ein Betrag eingehen, werden wir ihn eindeutig dem Vorgang zuordnen und Sie informieren.

Für Ihr privates Konto nennen unsere Zahlungsverkehrsbedingungen Montag bis Freitag, soweit kein gesetzlicher Feiertag am Sitz der Bank vorliegt, als Geschäftstage. Diese Information ändert keine gesetzlichen Ansprüche. Die Reklamation wurde am Dienstag, 13.10.2026, aufgenommen. Die technische Zusammenstellung ist noch nicht abgeschlossen. Unsere Ablehnung vom 14.10.2026 besteht nach dem aktuellen Bearbeitungsstand fort.

Diese Mitteilung enthält keine Erklärung einer Aufrechnung und beziffert keine gesonderte Ersatzforderung gegen Sie. Eine solche Erklärung ist bislang in diesem Vorgang nicht abgegeben worden. Die allein auf den Umgang mit den Zugangsdaten gestützte Ablehnung ist dem Schreiben vom 14.10. zu entnehmen.

Mit freundlichen Grüßen
Ottmar Wenzel, Reklamationsbearbeitung''')

MAILS = [
 (17,'Meldung_an_Bank','2026-10-13T09:12','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','Reklamation <reklamation@saaleufer-bank.example>','Vorgang SB-2026-1013-4821 Unterlagen zur fremden Zahlung', 'Sehr geehrte Damen und Herren,\n\nich hatte um 08.42 Uhr angerufen. Anbei finden Sie den Umsatzbeleg. Ich habe die 700 Euro nicht überwiesen. Ich wollte den Sessel verkaufen und kein Geld an einen fremden Empfänger schicken. Bitte erstatten Sie mir den Betrag und bestätigen Sie die Sperre. Das Passwort schicke ich selbstverständlich nicht mit.\n\nMit freundlichen Grüßen\nKunigunde Rabenstein',[7]),
 (18,'Bank_Eingangsbestätigung','2026-10-13T09:18','Reklamation <reklamation@saaleufer-bank.example>','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','Ihre Meldung ist erfasst', 'Sehr geehrte Frau Rabenstein,\n\nanbei erhalten Sie unsere Bestätigung zu Ihrer heutigen telefonischen Meldung. Ihre um 09.12 Uhr übersandten Unterlagen werden dem bereits angelegten Vorgang zugeordnet. Bitte beachten Sie: Die spätere E-Mail ist nicht der erstmalige Anzeigezeitpunkt. Für Rückfragen verwenden Sie die Vorgangsnummer.\n\nMit freundlichen Grüßen\nOttmar Wenzel',[10]),
 (19,'Kundin_an_Plattform','2026-10-13T10:04','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','Hofmarkt Sicherheit <sicherheit@hofmarkt.example>','Verdächtiger Kontakt zum Sessel HM-HAL-1210-85', 'Guten Tag,\n\nbitte sichern Sie meine Unterhaltung zu diesem Inserat. Nach dem Kontakt fehlen auf meinem Bankkonto 700 Euro. Der Sessel kostet 85 Euro, ist nicht abgeholt und es gab keine Kaufpreisgutschrift. Ich lege den Verlauf bei. Ich möchte den verdächtigen Link nicht noch einmal aufrufen. Bitte erklären Sie mir, ob von Ihrer Seite überhaupt eine Verkaufsbestätigung vorliegt.\n\nFreundliche Grüße\nKunigunde Rabenstein',[8,9]),
 (20,'Bank_Ablehnung_per_Mail','2026-10-14T14:10','Reklamation <reklamation@saaleufer-bank.example>','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','Entscheidung zu SB-2026-1013-4821', 'Sehr geehrte Frau Rabenstein,\n\nwir übersenden das Ergebnis unserer Reklamationsbearbeitung. Wegen der Eingabe Ihrer Zugangsdaten auf der fremden Seite sehen wir grobe Fahrlässigkeit und nehmen keine Erstattung vor. Die technische Zusammenstellung läuft daneben weiter. Einzelheiten entnehmen Sie bitte dem beigefügten Schreiben.\n\nMit freundlichen Grüßen\nOttmar Wenzel',[11]),
 (21,'Mandatsanfrage','2026-10-15T08:17','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','Daria Fenchel <daria.fenchel@kanzlei-fenchel.example>','Bitte um Hilfe Bank erstattet 700 Euro nicht', 'Sehr geehrte Frau Fenchel,\n\nich habe Ihre Kontaktdaten von meiner Schwester. Die Bank lehnt ab, weil ich auf die falsche Seite hereingefallen bin. Anbei das Schreiben. Ich schäme mich dafür, aber ich habe doch nichts überwiesen. Können Sie mir helfen und vorher sagen, was das ungefähr kostet? Ich kann morgen früh telefonieren. Meinen Sachverhalt habe ich am Dienstagmittag auch der Polizei übermittelt; eine Bestätigung habe ich noch nicht.\n\nMit freundlichen Grüßen\nKunigunde Rabenstein',[11]),
 (22,'Hofmarkt_Antwort','2026-10-15T11:22','Hofmarkt Sicherheit <sicherheit@hofmarkt.example>','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','HM-SICH-4821 Rückmeldung', 'Sehr geehrte Frau Rabenstein,\n\nunsere Rückmeldung finden Sie im Anhang. Bitte trennen Sie die Namen aus der Plattformunterhaltung und den Zahlungsempfänger; wir können die Identitäten nicht bestätigen. Die Daten sind unserem Sicherheitsvorgang zugeordnet. Bitte senden Sie keine Bankpasswörter.\n\nMit freundlichen Grüßen\nNele Dorn',[13]),
 (23,'Erinnerung_und_Datum','2026-10-16T09:38','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','Daria Fenchel <daria.fenchel@kanzlei-fenchel.example>','Meine Erinnerung und die Sache mit dem Montag', 'Sehr geehrte Frau Fenchel,\n\nder beigefügte Text stimmt nach unserem Gespräch. Ich hatte meiner Schwester erst vom Wochenende erzählt. Das war ungenau; es war Montagabend. Bei den Bildschirmen weiß ich nichts Sicheres mehr. Bitte schreiben Sie deshalb nicht, ich hätte irgendeine Überweisung bestätigt. Den Sessel habe ich weiterhin. Meine private Kopie der am Dienstag an die Polizei übermittelten Schilderung hängt ebenfalls an.\n\nMit freundlichen Grüßen\nKunigunde Rabenstein',[2,14]),
 (24,'Aktueller_Kontoauszug','2026-10-16T09:43','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','Daria Fenchel <daria.fenchel@kanzlei-fenchel.example>','Der angeforderte Auszug bis heute früh', 'Sehr geehrte Frau Fenchel,\n\nhier der Auszug. Die 12,80 Euro sind aus der Apotheke; das ist keine Teilzahlung der Bank für den Betrug. Ich wollte das gleich dazusagen. Die normalen Lastschriften laufen weiter. Eine Gutschrift über 700 Euro ist noch nicht da. Bitte nehmen Sie nicht den ganzen Kontosaldo als Schaden.\n\nMit freundlichen Grüßen\nKunigunde Rabenstein',[12]),
 (25,'Rueckrufstatus_Bank','2026-10-16T11:34','Reklamation <reklamation@saaleufer-bank.example>','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','SB-2026-1013-4821 Stand der Rückholung', 'Sehr geehrte Frau Rabenstein,\n\nanbei der aktuelle Stand. Die Rückholanfrage ist weiter offen. Es gibt noch keinen Zahlungseingang. Die Rechtsabteilung ist über die Anfrage Ihrer Anwältin informiert; ein eigenes Antwortschreiben liegt noch nicht vor.\n\nMit freundlichen Grüßen\nOttmar Wenzel',[16]),
 (26,'Entwuerfe_zur_Ruecksprache','2026-10-16T13:40','Daria Fenchel <daria.fenchel@kanzlei-fenchel.example>','Kunigunde Rabenstein <kunigunde.rabenstein@postfach.example>','DF 166/26 Bitte prüfen Sie die Sachverhaltsschilderung', 'Sehr geehrte Frau Rabenstein,\n\nanbei erhalten Sie den Arbeitsentwurf des Bankbriefs und unsere Betragsübersicht. Bitte prüfen Sie die Sachverhaltsschilderung. Die rechtliche Schlusskontrolle und die Prüfung der Anlagen bearbeitet die Kanzlei. Das Schreiben wurde noch nicht an die Bank versandt. Eine Klage ist vorbereitet, aber weder eingereicht noch zur Einreichung freigegeben.\n\nBitte informieren Sie uns sofort über neue Buchungen oder eine Antwort der Bank. Eine mögliche spätere Gegenforderung ist in der Tabelle nicht mit der bislang fehlenden Erstattung verrechnet.\n\nMit freundlichen Grüßen\nDaria Fenchel, Rechtsanwältin',[4,28]),
]

XLSX_JS = r'''
import fs from 'node:fs/promises';
import path from 'node:path';
import {Workbook,SpreadsheetFile} from '@oai/artifact-tool';
const [folder,qa]=process.argv.slice(2),reports=[];
const serial=s=>Date.parse(s+'T00:00:00Z')/86400000+25569;
function sheet(wb,name,title,headers,widths,rows){
 const w=wb.worksheets.add(name),last=String.fromCharCode(64+headers.length);w.showGridLines=false;
 w.getRange(`A1:${last}${rows.length+6}`).format.font={name:'Arial',size:11,color:'#172B3A'};
 w.getRange('A2').values=[[title]];w.getRange('A2').format.font={name:'Arial',size:15,bold:true};w.getRange('A2').format.rowHeight=28;
 w.getRange(`A5:${last}5`).values=[headers];w.getRange(`A5:${last}5`).format.fill='#334B62';w.getRange(`A5:${last}5`).format.font={name:'Arial',size:11,color:'#FFFFFF',bold:true};w.getRange(`A5:${last}5`).format.rowHeight=30;
 w.getRange(`A5:${last}${rows.length+5}`).format.wrapText=true;w.getRange(`A5:${last}${rows.length+5}`).format.verticalAlignment='center';
 widths.forEach((x,i)=>w.getRange(`${String.fromCharCode(65+i)}:${String.fromCharCode(65+i)}`).format.columnWidth=x);
 w.getRange(`A6:${last}${rows.length+5}`).values=rows;w.getRange(`A6:${last}${rows.length+5}`).format.rowHeight=54;
 rows.forEach((_,i)=>{if(i%2)w.getRange(`A${i+6}:${last}${i+6}`).format.fill='#F1F4F6';});w.freezePanes.freezeRows(5);return w;
}
for(let kind=0;kind<2;kind++){
 const wb=Workbook.create(),views=[],checks=[];let input,output,expectedChange;
 if(kind===0){
 const events=[['2026-10-12','17.58','Inserat eingestellt','08_Inserat_Hofmarkt.pdf'],['2026-10-12','19.03','Nachricht zur angeblichen Bestätigung','09_Nachrichtenverlauf_Hofmarkt.pdf'],['2026-10-12','19.17','700 EUR belastet; Systemzeit, keine Autorisierungsfeststellung','07_Belastungsbeleg_12_Oktober.pdf'],['2026-10-13','08.42','Erste Anzeige bei Bank telefonisch','10_Eingang_Bank_und_Sperrung.pdf'],['2026-10-13','08.51','Onlinezugang gesperrt','10_Eingang_Bank_und_Sperrung.pdf'],['2026-10-14','14.10','Ablehnung per E-Mail übersandt','20_Bank_Ablehnung_per_Mail.eml'],['2026-10-15','11.22','Plattform antwortet','22_Hofmarkt_Antwort.eml'],['2026-10-16','08.00','Auszugsstand ohne Erstattung','12_Folgeauszug_bis_16_Oktober.pdf'],['2026-10-16','09.10','Mandatsaufnahme','01_Mandatsauftrag_16_Oktober.docx'],['2026-10-16','10.05','Banktelefonat der Anwältin','03_Telefonnotiz_Bankkontakt.docx']];
 let w=sheet(wb,'Chronologie','Kunigunde Rabenstein Ereignisse',['Datum','Uhrzeit','Ereignis','Beleg'],[18,12,70,52],events.map(x=>[serial(x[0]),(Number(x[1].split('.')[0])*60+Number(x[1].split('.')[1]))/1440,...x.slice(2)]));w.getRange('A6:A15').setNumberFormat('yyyy-mm-dd');w.getRange('B6:B15').setNumberFormat('hh:mm');w.getRange('A6:B15').format.horizontalAlignment='center';views.push([w,15,4]);
 w=sheet(wb,'Anlagen','Zuordnung zum Bankbrief und Klageentwurf',['Kennung','Datei','Gegenstand','Stand'],[13,53,40,38],[['K1','07_Belastungsbeleg_12_Oktober.pdf','Belastung und Wertstellung','vorhanden'],['K2','08_Inserat_Hofmarkt.pdf','Sessel für 85 EUR','vorhanden'],['K3','09_Nachrichtenverlauf_Hofmarkt.pdf','Verlauf; keine echten Identitäten','vorhanden'],['K4','10_Eingang_Bank_und_Sperrung.pdf','Anzeige 13.10., 08.42 Uhr','vorhanden'],['K5','11_Ablehnung_Saaleufer_Bank.pdf','Ablehnung 14.10.','vorhanden'],['K6','12_Folgeauszug_bis_16_Oktober.pdf','Keine Erstattung bis 16.10. 08.00','vorhanden'],['K7','02_Gedaechtnisprotokoll_Rabenstein.docx','Erinnerung, per E-Mail 23 bestätigt','vorhanden'],['K8','06_Vollmacht_Rabenstein.docx','Vertretungsumfang','vorhanden'],['offen','Bankprotokolle noch nicht übergeben','Gerät, Authentifizierung, Freigabe','angefordert'],['offen','Vollständiger Rahmenvertrag fehlt','Kontovertrag, keine Ergänzung aus Erinnerung','angefordert']]);views.push([w,15,4]);
 w=sheet(wb,'Zeitabstand','Kalendertage seit der ersten Meldung',['Feld','Wert','Bedeutung'],[28,25,92],[['Stand',serial('2026-10-16'),'Fester Aktenstand, keine aktuelle Uhr'],['Erste Anzeige',serial('2026-10-13'),'Telefonisch 08.42 Uhr; spätere E-Mail ist Ergänzung'],['Kalendertage',null,'Reine Datendifferenz; keine gesetzliche Fristberechnung'],['Unterlagenstatus','offen','Bankprotokolle und vollständiger Vertrag fehlen.'],['Benutzung','Datum ändern','Datum als echten Excel-Wert eintragen; ursprünglichen Beleg behalten.']]);w.getRange('B6:B7').setNumberFormat('yyyy-mm-dd');w.getRange('B6:B8').format.horizontalAlignment='center';w.getRange('B8').formulas=[['=B6-B7']];checks.push({sheet:w.name,cell:'B8',expected:3});views.push([w,10,3]);input=[w,'B6'];output=[w,'B8'];expectedChange=4;
 }else{
 let w=sheet(wb,'Betragsstand','Streitige Zahlung und tatsächliche Buchungen',['Position','Betrag EUR','Herkunft oder Bedeutung'],[47,20,94],[['Streitige Belastung',700,'Beleg 07, ZV-121026-1917; positiver Anspruchsbetrag'],['Erstattet durch Bank',0,'Auszug 12 bis 16.10. 08.00: keine passende Gutschrift'],['Noch nicht erstattet',null,'Belastung abzüglich tatsächlich zugeordneter Erstattung'],['Sonstige Schäden',null,'Nicht beziffert, nicht als Null oder automatischer Anspruch behandelt'],['Späterer Ersatzanspruch Bank',null,'Nicht beziffert; rechtlich und tatsächlich gesondert zu prüfen'],['Anzeige Bank','13.10.2026 08.42','Beleg 10; kein Formelersatz für die Fristprüfung'],['Stand','16.10.2026 08.00','Spätere Buchungen müssen nachgetragen werden.']]);w.getRange('B6:B10').setNumberFormat('#,##0.00');w.getRange('B6:B10').format.horizontalAlignment='center';w.getRange('B8').formulas=[['=B6-B7']];checks.push({sheet:w.name,cell:'B8',expected:700});views.push([w,12,3]);input=[w,'B7'];output=[w,'B8'];expectedChange=699;
 w=sheet(wb,'Kontoumsaetze','Kontrollrechnung zum Folgeauszug',['Datum','Umsatz','Betrag EUR','Saldo EUR'],[18,66,19,20],[['2026-10-12','Anfangssaldo vor streitiger Belastung',null,2146.38],['2026-10-12','Überweisung N. Kramer',-700,null],['2026-10-13','Stadtwerke Abschlag',-62,null],['2026-10-14','Lebensmittel',-27.45,null],['2026-10-15','Apothekenerstattung, anderer Vorgang',12.8,null]].map(x=>[serial(x[0]),...x.slice(1)]));w.getRange('A6:A10').setNumberFormat('yyyy-mm-dd');w.getRange('A6:A10').format.horizontalAlignment='center';w.getRange('C6:D10').setNumberFormat('#,##0.00');for(let r=7;r<=10;r++)w.getRange('D'+r).formulas=[['=D'+(r-1)+'+C'+r]];checks.push({sheet:w.name,cell:'D10',expected:1369.73});views.push([w,10,4]);
 w=sheet(wb,'Bearbeitung','Getrennte Bearbeitungsfragen',['Gegenstand','Stand','Nächste Arbeit'],[38,28,91],[['Erstattung 700 EUR','verlangt, offen','Aktuellen Auszug prüfen; Entwurf nach Quellenprüfung freigeben.'],['Rückholung Empfängerkonto','offen','Kein Eingang bestätigt; nicht als Erstattung verbuchen.'],['Bank behauptet grobe Fahrlässigkeit','streitig','Vorwurf und Nachweise gesondert auswerten; kein Anerkenntnis.'],['Eigenes betrügerisches Handeln Kundin','nicht behauptet','Telefonnotiz prüfen; schriftliche Antwort noch angefordert.'],['Späterer Bankregress','nicht beziffert','Keinen Betrag erfinden und nicht mit der Erstforderung verrechnen.'],['Weitere Schäden oder Zinsen','offen','Nur mit Belegen, Voraussetzungen und zeitlicher Zuordnung ergänzen.'],['Versand und Einreichung','nicht freigegeben','Daria Fenchel prüft Text und Anlagen; Mandantin entscheidet über Klage.']]);views.push([w,12,3]);
 }
 wb.recalculate();for(const x of checks){x.actual=wb.worksheets.getItem(x.sheet).getRange(x.cell).values[0][0];if(Math.abs(x.actual-x.expected)>1e-7)throw Error(JSON.stringify(x));}
 const old=input[0].getRange(input[1]).values[0][0];input[0].getRange(input[1]).values=[[old+1]];wb.recalculate();const changed=output[0].getRange(output[1]).values[0][0];if(changed!==expectedChange)throw Error('Rechenprobe fehlgeschlagen');input[0].getRange(input[1]).values=[[old]];wb.recalculate();
 const errors=await wb.inspect({kind:'match',searchTerm:'#REF!|#DIV/0!|#VALUE!|#NAME\\?|#N/A|#NUM!|#NULL!|#SPILL!|#CALC!',options:{useRegex:true,maxResults:30},maxChars:2000});
 const file=kind===0?'27_Chronologie_und_Anlagen.xlsx':'28_Betragsrechnung_und_Status.xlsx';const renderdir=path.join(qa,file.replace('.xlsx',''));await fs.mkdir(renderdir,{recursive:true});
 for(const [w,rows,cols] of views){const b=await wb.render({sheetName:w.name,range:`A1:${String.fromCharCode(64+cols)}${rows}`,scale:1,format:'png'});await fs.writeFile(path.join(renderdir,w.name+'.png'),new Uint8Array(await b.arrayBuffer()));}
 const result=await SpreadsheetFile.exportXlsx(wb);await result.save(path.join(folder,file));await fs.rm(path.join(folder,file+'.inspect.ndjson'),{force:true});reports.push({file,sheets:views.map(([w,rows,cols])=>({name:w.name,range:`A1:${String.fromCharCode(64+cols)}${rows}`})),formula_checks:checks,input_change_test:{cell:input[1],changed,restored:true},errors:errors.ndjson});
}
await fs.writeFile(path.join(qa,'xlsx-pruefung.json'),JSON.stringify(reports,null,2)+'\n');
'''

def write_docs():
    from docx import Document
    from docx.shared import Mm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
    fonts=Path('/System/Library/Fonts/Supplemental')
    for key,name in [('Case','Times New Roman.ttf'),('CaseBold','Times New Roman Bold.ttf')]:
        pdfmetrics.registerFont(TTFont(key,str(fonts/name)))
    body=ParagraphStyle('body',fontName='Case',fontSize=11,leading=14.2,spaceAfter=10)
    head=ParagraphStyle('head',parent=body,fontName='CaseBold',spaceBefore=8,spaceAfter=10,keepWithNext=True)
    title=ParagraphStyle('title',parent=head,fontSize=16,leading=19,spaceAfter=18)
    for item in DOCS:
        p=OUT/item['file']
        if p.suffix=='.docx':
            d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(22);s.left_margin=s.right_margin=Mm(23)
            for name in ['Normal','Title','Heading 1']:
                st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
                st.paragraph_format.space_after=Pt(9);st.paragraph_format.line_spacing=1.08
                for attr in list(st.element.rPr.rFonts.attrib):
                    if attr.endswith('Theme') or attr.endswith('theme'):del st.element.rPr.rFonts.attrib[attr]
                for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
            d.styles['Title'].font.size=Pt(16);d.styles['Title'].font.bold=True;d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
            d.add_paragraph(item['title'],'Title')
            for text in item['paragraphs']:d.add_paragraph(text,'Heading 1' if re.match(r'^\d+\. ',text) and len(text)<100 else 'Normal')
            footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run('DF 166/26  |  ').font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
            d.core_properties.author='Kanzlei Fenchel';d.core_properties.title=item['title'];d.core_properties.language='de-DE';d.core_properties.created=dt.datetime(2026,10,16);d.core_properties.modified=dt.datetime(2026,10,16);d.save(p)
        else:
            flow=[Paragraph(escape(item['title']),title)]
            for text in item['paragraphs']:
                if item['file'].startswith('12_') and re.match(r'^1[2345]\.10\.2026, ',text):continue
                flow.append(Paragraph(escape(text).replace('\n','<br/>'),head if re.match(r'^\d+\. ',text) and len(text)<100 else body))
                if item['file'].startswith('12_') and text.startswith('Stand vor der streitigen'):
                    rows=[['Datum','Buchung','Betrag EUR','Saldo EUR'],['12.10.2026','N. Kramer','−700,00','1.446,38'],['13.10.2026','Stadtwerke','−62,00','1.384,38'],['14.10.2026','Lebensmittel','−27,45','1.356,93'],['15.10.2026','Apotheke','12,80','1.369,73']]
                    t=Table(rows,colWidths=[74,183,90,90],repeatRows=1)
                    t.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'Case'),('FONTSIZE',(0,0),(-1,-1),11),('FONTNAME',(0,0),(-1,0),'CaseBold'),('BACKGROUND',(0,0),(-1,0),'#E9EDF0'),('GRID',(0,0),(-1,-1),0.4,'#D9D9D9'),('ALIGN',(2,0),(-1,-1),'RIGHT'),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)]));flow.extend([t,Spacer(1,12)])
            def footer(canvas,document):
                canvas.setFont('Case',9);canvas.drawRightString(530,29,f'DF 166/26  |  {document.page}')
            SimpleDocTemplate(str(p),pagesize=A4,rightMargin=65,leftMargin=65,topMargin=58,bottomMargin=48,title=item['title'],author='Fiktive Beteiligte der Akte Rabenstein').build(flow,onFirstPage=footer,onLaterPages=footer)

def write_mails():
    mapping={int(x['file'][:2]):x['file'] for x in DOCS}
    mapping.update({27:'27_Chronologie_und_Anlagen.xlsx',28:'28_Betragsrechnung_und_Status.xlsx'})
    for n,name,date,sender,to,subject,body,attachments in MAILS:
        m=EmailMessage(policy=SMTP);m['From']=sender;m['To']=to;m['Subject']=subject;m['Date']=format_datetime(dt.datetime.fromisoformat(date).replace(tzinfo=ZoneInfo('Europe/Berlin')));m['Message-ID']=f'<rabenstein-{n}-20261016@akten.example>'
        if n in [18,20,25]:m['In-Reply-To']='<rabenstein-17-20261016@akten.example>'
        m.set_content(body)
        for k in attachments:
            p=OUT/mapping[k];sub={'.pdf':'pdf','.docx':'vnd.openxmlformats-officedocument.wordprocessingml.document','.xlsx':'vnd.openxmlformats-officedocument.spreadsheetml.sheet'}[p.suffix]
            m.add_attachment(p.read_bytes(),maintype='application',subtype=sub,filename=p.name)
        m.set_boundary(f'case-rabenstein-{n}-20261016')
        (OUT/f'{n:02d}_{name}.eml').write_bytes(m.as_bytes())

def write_readme():
    (OUT/'README.md').write_text(f'''# Kunigunde Rabenstein und die nicht erstattete Zahlung

Diese deutsche Arbeitsakte eignet sich für die anwaltliche Bearbeitung eines Phishing-Vorfalls beim privaten Onlineverkauf. Sie enthält die Unterlagen einer Kundin aus Halle (Saale), die eine Belastung über 700 Euro bestreitet und von ihrer Bank allein wegen behaupteter grober Fahrlässigkeit keine Erstattung erhält. Bankbrief und Klageentwurf sind bearbeitbare Arbeitsstände.

## 1. Benutzung

Lesen Sie zuerst den Mandatsauftrag, dann die Originalbelege und den Schriftverkehr. Arbeiten Sie die belegte Chronologie, den Zahlungsauftrag, die Erstattung und eine etwaige spätere Ersatzforderung getrennt heraus. Ergänzen Sie die Entwürfe nur anhand überprüfter Tatsachen und Quellen. Die Erstattungsforderung soll nicht mit dem gesamten Kontosaldo, dem Kaufpreis des Sessels oder der Apothekenerstattung verwechselt werden.

Die Akte folgt dem im amtlichen Urteil des EuGH vom 15.10.2026, C-70/25, Tukowiecka, Rn. 14 bis 20, beschriebenen Kernablauf. Das Ausgangsverfahren betraf 3.000 PLN, etwa 705 Euro; der deutsche Übungsfall verwendet 700 Euro. Ort, Personen, Beträge in den sonstigen Kontoumsätzen, Korrespondenz und Daten sind für das deutsche Szenario erfunden. Sie werden nicht als zusätzliche Tatsachen des EuGH-Verfahrens ausgegeben. Der Kaufpreis des Gegenstands und der streitige Zahlungsbetrag sind verschiedene Größen. Die Kundin berichtet die Weitergabe eines SMS-Codes. Dessen Inhalt und der gesamte SCA- oder Freigabeablauf sind nicht festgestellt.

## 2. Aktenbestand

Die 28 Originaldateien umfassen sechs Word-Dokumente, zehn PDFs, zehn E-Mails mit tatsächlichen Anhängen und zwei Excel-Arbeitsmappen. Die Arbeitsmappen enthalten sechs erläuterte Tabellenblätter, berechnete Salden, eine getrennte Erstattungsübersicht und die Zuordnung der Anlagen K1 bis K8. Eine nicht bekannte Ersatzforderung bleibt leer; sie wird nicht mit null angenommen oder gegen die Erstattung verrechnet.

Alle Personen, Unternehmen, Kontokennungen und Kontaktanschriften sind erfunden. Die postalischen Anschriften sind als fiktiv gekennzeichnet; E-Mail-Adressen verwenden reservierte `.example`-Domains. Keine Datei enthält einen aktiven Phishing-Link, ein Passwort, eine nutzbare IBAN oder ein amtliches Polizeischreiben. Die Strafanzeige ist als private Kopie der eigenen Schilderung enthalten, ohne nachgeahmte amtliche Bestätigung.

<!-- reserved-example-contacts -->

<!-- BEGIN gesamt-pdf-section (autogen) -->
## 3. Downloads

{NOTICE}

| Format | Download |
| --- | --- |
| Gesamt-PDF | [Gesamtakte öffnen](./gesamt-pdf/{SLUG}_gesamt.pdf) |
| Originaldateien | [Akten-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/phishing-v445.35.4/testakte-{SLUG}.zip) |
| Einzel-PDFs | [Einzel-PDF-ZIP](https://github.com/Klotzkette/claude-fuer-deutsches-recht/releases/download/phishing-v445.35.4/testakte-{SLUG}-einzelpdfs.zip) |

<!-- END gesamt-pdf-section (autogen) -->

## 4. Weiterarbeiten

Nutzen Sie den [Phishing-Vorfall-Prüfer](../../phishing-vorfall-pruefer/README.md). Der Aktenstand ist der 16.10.2026. Die Anwältin hat das Bankanschreiben und die Klage nicht versandt oder eingereicht. Eine Musterlösung gehört nicht zum Mandatsdossier; Quellenprüfung und getrennte Prüfkriterien liegen außerhalb der Akte.


[Repository-Übersicht](../../README.md) · [Testaktenübersicht](../README.md) · [Download-Index](../../ASSET_INDEX.md)
''',encoding='utf-8')

def manifest():
    from pypdf import PdfReader
    from openpyxl import load_workbook
    files=[];mails=[]
    for p in sorted(OUT.iterdir()):
        if p.suffix.lower() not in ['.docx','.pdf','.eml','.xlsx']:continue
        x=dict(file=p.name,type=p.suffix[1:],bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest())
        if p.suffix=='.pdf':x['pages']=len(PdfReader(p).pages)
        if p.suffix=='.eml':
            msg=BytesParser(policy=SMTP).parsebytes(p.read_bytes());a=[]
            for part in msg.iter_attachments():
                fn=part.get_filename();data=part.get_payload(decode=True);assert data==(OUT/fn).read_bytes();a.append(dict(file=fn,sha256=hashlib.sha256(data).hexdigest()))
            mails.append(dict(file=p.name,message_id=str(msg['Message-ID']),attachments=a))
        files.append(x)
    assert len(files)==28 and len(mails)==10
    report=dict(slug=SLUG,stand='2026-10-16',files=files,mails=mails,workbooks=json.loads((QA/'xlsx-pruefung.json').read_text()),annex_map={f'K{i}':name for i,name in enumerate(['07_Belastungsbeleg_12_Oktober.pdf','08_Inserat_Hofmarkt.pdf','09_Nachrichtenverlauf_Hofmarkt.pdf','10_Eingang_Bank_und_Sperrung.pdf','11_Ablehnung_Saaleufer_Bank.pdf','12_Folgeauszug_bis_16_Oktober.pdf','02_Gedaechtnisprotokoll_Rabenstein.docx','06_Vollmacht_Rabenstein.docx'],1)})
    for record in report['workbooks']:
        w=load_workbook(OUT/record['file'],data_only=True)
        for c in record['formula_checks']:assert abs(w[c['sheet']][c['cell']].value-c['expected'])<1e-7
    target=ROOT/'quality/phishing-vorfall-pruefer/testakte-pruefung.json';target.parent.mkdir(parents=True,exist_ok=True);target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')

def main():
    p=argparse.ArgumentParser();p.add_argument('--documents-only',action='store_true');p.add_argument('--manifest-only',action='store_true');a=p.parse_args()
    OUT.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True);contents()
    if a.manifest_only:manifest();return
    write_docs();write_readme()
    if a.documents_only:return
    target=QA/'node_modules'
    if not target.exists():target.symlink_to(RUNTIME/'node/node_modules',target_is_directory=True)
    (QA/'workbooks.mjs').write_text(XLSX_JS)
    subprocess.run([str(RUNTIME/'node/bin/node'),str(QA/'workbooks.mjs'),str(OUT),str(QA)],check=True)
    write_mails();manifest()
    print('28 Originaldateien erzeugt; E-Mail-Anhänge und gespeicherte Formelwerte geprüft.')

if __name__=='__main__':main()
