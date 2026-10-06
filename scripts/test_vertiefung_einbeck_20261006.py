#!/usr/bin/env python3
"""Rechen-, Inhalts- und Paketprüfung der Fortschreibung. Keine QA in der Akte."""
from pathlib import Path
import os,sys,json,io,zipfile,shutil,subprocess,xml.etree.ElementTree as ET,hashlib,importlib.util,re
from openpyxl import load_workbook
from pypdf import PdfReader
from office_process import run_office

ROOT=Path(__file__).resolve().parents[1];CASE=ROOT/'testakten/bauwirtschaft-hoai-buergerhaus-einbeck'
QA=Path(os.environ.get('EINBECK_VERTIEFUNG_QA','/tmp/bauwirtschaft-vertiefung-20261006/einbeck'))
N='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
DC='{http://purl.org/dc/elements/1.1/}';CP='{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}'

def patch(path,ranges=None,mutation=None):
    out=io.BytesIO()
    with zipfile.ZipFile(path) as z,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as o:
        for inf in z.infolist():
            data=z.read(inf)
            if inf.filename=='docProps/core.xml':
                r=ET.fromstring(data)
                for tag in (DC+'creator',CP+'lastModifiedBy'):
                    e=r.find(tag)
                    if e is None:e=ET.SubElement(r,tag)
                    e.text='Klotzkette'
                data=ET.tostring(r,encoding='utf-8',xml_declaration=True)
            elif inf.filename=='xl/workbook.xml':
                r=ET.fromstring(data);names=r.find(N+'definedNames')
                if names is None:names=ET.SubElement(r,N+'definedNames')
                if ranges:
                    for e in list(names):
                        if e.get('name') in ('_xlnm.Print_Area','_xlnm.Print_Titles'):names.remove(e)
                    for i,(name,area) in enumerate(ranges):
                        e=ET.SubElement(names,N+'definedName',name='_xlnm.Print_Area',localSheetId=str(i));e.text=f"'{name}'!{area}"
                        e=ET.SubElement(names,N+'definedName',name='_xlnm.Print_Titles',localSheetId=str(i));e.text=f"'{name}'!$7:$7"
                calc=r.find(N+'calcPr')
                if calc is None:calc=ET.SubElement(r,N+'calcPr')
                calc.attrib.update(calcMode='auto',fullCalcOnLoad='1',forceFullCalc='1')
                data=ET.tostring(r,encoding='utf-8',xml_declaration=True)
            elif inf.filename.startswith('xl/worksheets/sheet') and inf.filename.endswith('.xml'):
                r=ET.fromstring(data)
                if ranges:
                    setup=r.find(N+'pageSetup')
                    if setup is None:setup=ET.SubElement(r,N+'pageSetup')
                    setup.attrib.update(paperSize='9',orientation='landscape',fitToWidth='1',fitToHeight='0')
                    prop=r.find(N+'sheetPr')
                    if prop is None:prop=ET.Element(N+'sheetPr');r.insert(0,prop)
                    ps=prop.find(N+'pageSetUpPr')
                    if ps is None:ps=ET.SubElement(prop,N+'pageSetUpPr')
                    ps.set('fitToPage','1')
                    margins=r.find(N+'pageMargins')
                    if margins is None:margins=ET.SubElement(r,N+'pageMargins')
                    margins.attrib.update(left='.25',right='.25',top='.35',bottom='.35',header='.15',footer='.15')
                if mutation:
                    for change in mutation:
                        if inf.filename!=f"xl/worksheets/sheet{change['sheet']}.xml":continue
                        e=r.find(f".//{N}c[@r='{change['cell']}']")
                        if e is None:
                            rownum=''.join(filter(str.isdigit,change['cell']));row=r.find(f".//{N}row[@r='{rownum}']")
                            e=ET.SubElement(row,N+'c',r=change['cell'])
                        for child in list(e):e.remove(child)
                        e.attrib.pop('t',None);v=change['value']
                        if isinstance(v,str):e.set('t','inlineStr');ET.SubElement(ET.SubElement(e,N+'is'),N+'t').text=v
                        elif v is not None:ET.SubElement(e,N+'v').text=str(v)
                    for e in r.findall(f'.//{N}c'):
                        if e.find(N+'f') is not None:
                            for v in e.findall(N+'v'):e.remove(v)
                data=ET.tostring(r,encoding='utf-8',xml_declaration=True)
            o.writestr(inf,data)
    path.write_bytes(out.getvalue())

def formulas(p):
    w=load_workbook(p,data_only=False);d={(s.title,c.coordinate):re.sub(r"'(\w+)'!",r'\1!',c.value) for s in w for row in s for c in row if c.data_type=='f'};w.close();return d

def native(paths,label):
    dest=QA/'native'/label;dest.mkdir(parents=True,exist_ok=True)
    binary=os.environ['SOFFICE']
    assert 'codex-primary-runtime/dependencies/bin/override/soffice' in binary
    cmd=[binary,f"-env:UserInstallation={(dest/'profile').as_uri()}",'--headless','--convert-to','xlsx','--outdir',str(dest),*[str(x) for x in paths]]
    code,logs=run_office(cmd,timeout=180,env=os.environ.copy());(dest/'lauf.log').write_text(logs)
    assert code==0,(code,logs)
    for p in paths:assert (dest/p.name).exists(),logs
    return [dest/p.name for p in paths]

def check_value(path,sheet,cell,expected):
    w=load_workbook(path,data_only=True);actual=w[sheet][cell].value;w.close()
    if isinstance(expected,(int,float)):assert isinstance(actual,(int,float)) and abs(actual-expected)<.001,(path.name,sheet,cell,actual,expected)
    else:assert actual==expected,(path.name,sheet,cell,actual,expected)
    return actual

def package_checks():
    """Wendet unveränderte zentrale Prüffunktionen nur auf diesen Fall an."""
    from testakte_disclaimer import NOTICE_BYTES
    from testakte_file_filter import include_in_working_dump
    dist=Path(os.environ.get('EINBECK_VERTIEFUNG_DIST','/tmp/bauwirtschaft-vertiefung-20261006/dist'))
    def mod(name,file):
        sp=importlib.util.spec_from_file_location(name,ROOT/'scripts'/file);m=importlib.util.module_from_spec(sp);sp.loader.exec_module(m);return m
    z=mod('einbeck_zip_validator','validate-testakten-release-zips.py');e=mod('einbeck_pdfzip_validator','validate-testakten-einzelpdf-zips.py')
    r=mod('einbeck_readme_validator','validate-testakten-readme-downloads.py');g=mod('einbeck_total_validator','validate-testakten-gesamt-pdf.py')
    originals=[p for p in CASE.iterdir() if include_in_working_dump(p,CASE)];assert len(originals)==55
    native=dist/f'testakte-{CASE.name}.zip';pdfzip=dist/f'testakte-{CASE.name}-einzelpdfs.zip'
    z.assert_same(native.name,z.expected_entries(CASE),z.zip_entries(native,require_notice=True))
    e.assert_same(pdfzip.name,['README.txt',*e.expected_arcnames(CASE)],e.zip_entries(pdfzip,expected_suffix='.pdf'))
    with zipfile.ZipFile(native) as archive:
        for p in originals:assert archive.read(p.name)==p.read_bytes(),p.name
    total=CASE/'gesamt-pdf'/f'{CASE.name}_gesamt.pdf';assert not g.pdf_content_errors(total.read_bytes())
    errors=[];r.validate_local_readme(CASE.name,CASE,errors);assert not errors,errors
    text='\n'.join(p.extract_text() for p in PdfReader(total).pages)
    needles=['Nachlieferung der Projektunterlagen','Erläuterung des Aufmaßes AM 07','Erläuterung der Gerätebereitschaft N01 5','Aufmaß und Schlussrechnung Los 1','vielen Dank für Ihre Erläuterung','Beobachtungen an Rinne und Saalwand','Angebot zur Untersuchung des Rinnenanschlusses','Leistungsnachweise und offene Unterlagen','Mängel und Befundverfolgung','Honorarzuordnung und Leistungsanteile','Herr Feldmann und ich haben Ihre Ergänzungen']
    for needle in needles:assert needle in text,needle
    assert 'Diese Testakte wurde' not in text and 'This test case file' not in text
    new=[p for p in originals if int(p.name[:2])>=45];counts={}
    with zipfile.ZipFile(pdfzip) as archive:
        for p in new:counts[p.name]=len(PdfReader(io.BytesIO(archive.read(p.stem+'.pdf'))).pages)
    result={'originals':55,'new_originals':11,'individual_pdfs':55,'new_pdf_pages':sum(counts.values()),'gesamt_pages':len(PdfReader(total).pages),'new_pages_by_file':counts,'zip_bytes_identical':True,'central_functions':'Originalformat-ZIP, Einzel-PDF-ZIP, Gesamt-PDF-Warntextprüfung und lokale README bestanden'}
    (QA/'paket-pruefung.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))

def main():
    reports=json.loads((QA/'tabellen-pruefung.json').read_text());paths=[];fm={}
    for r in reports:
        p=CASE/r['file'];patch(p,r['ranges']);paths.append(p);fm[p.name]=formulas(p)
    outs=native(paths,'basis');checks=[]
    expected={48:[('Abgleich','E26',115393.8),('Abgleich','F26',114153.8),('Abgleich','F30',52543.02),('Abgleich','F32','Zahlungsbeleg fehlt')],53:[('Vorgänge','B5',3),('Beobachtungen','C16',47),('Beobachtungen','D16',58)],54:[('Honorar','C22',54000),('Honorar','E22',53040),('Honorar','F22',960),('Honorar','G26',3)]}
    for r,p in zip(reports,outs):
        assert fm[p.name]==formulas(p),'Native Neuberechnung verändert Formeln'
        for s,c,v in expected[r['n']]:check_value(p,s,c,v)
        w=load_workbook(p,data_only=True)
        assert not [(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e']
        w.close();shutil.copyfile(p,CASE/p.name);patch(CASE/p.name,r['ranges'])
        checks.append({'file':p.name,'formulas':len(fm[p.name]),'baseline':'ok'})
    probein=QA/'mutationen';probein.mkdir(exist_ok=True);probes=[]
    for r in reports:
        source=CASE/r['file'];w=load_workbook(source);sheets=w.sheetnames;w.close()
        for pr in r['probes']:
            p=probein/f"probe-{len(probes):02}.xlsx";shutil.copyfile(source,p)
            patch(p,mutation=[dict(sheet=sheets.index(pr['sheet'])+1,cell=pr['cell'],value=pr['value'])]);probes.append((p,pr))
    source=CASE/reports[0]['file']
    extras=[([dict(sheet=3,cell='B12',value='Belegt'),dict(sheet=3,cell='B13',value=0)],'Zahlungsbeleg fehlt'),([dict(sheet=3,cell='B12',value='Belegt'),dict(sheet=3,cell='B13',value=0),dict(sheet=3,cell='B14',value='Zahlungsabgleich 06.10.2026')],52543.02)]
    for mut,expected in extras:
        p=probein/f"probe-{len(probes):02}.xlsx";shutil.copyfile(source,p);patch(p,mutation=mut);probes.append((p,dict(outSheet='Abgleich',outCell='F32',expected=expected,mutation=mut)))
    outputs=native([p for p,pr in probes],'mutation');result=[]
    for p,(_,pr) in zip(outputs,probes):result.append({**pr,'actual':check_value(p,pr['outSheet'],pr['outCell'],pr['expected'])})
    (QA/'native-pruefung.json').write_text(json.dumps({'baselines':checks,'mutation_tests':result},ensure_ascii=False,indent=2))
    baseline=json.loads((QA/'baseline-originale.json').read_text())
    assert all(hashlib.sha256((CASE/n).read_bytes()).hexdigest()==h for n,h in baseline.items())
    print(json.dumps({'originale_unveraendert':len(baseline),'xlsx':checks,'mutation_tests':len(result)},ensure_ascii=False))

if __name__=='__main__':
    if '--packages-only' in sys.argv:package_checks()
    else:main()
