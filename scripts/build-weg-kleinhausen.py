#!/usr/bin/env python3
"""Baut die einfache WEG-Akte Kleinhausen mit dem gemeinsamen Office-Builder.

--qa-dir ist Pflicht. Gesamt-PDFs und Archive entstehen mit den zentralen
Exportwerkzeugen. Die vorhandenen vier WEG-Akten werden nicht angefasst.
"""
from pathlib import Path
import importlib.util
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('kleinhausen_native_builder', ROOT / 'scripts/build-berlin-bildungsakten.py')
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)
from weg_kleinhausen_daten import CASES

base_word = builder.word


def word(item, target):
    base_word(item, target)
    if not item.get('tables'):
        return
    doc = Document(target)
    for table in doc.tables:
        props = table._tbl.tblPr
        borders = props.find(qn('w:tblBorders'))
        if borders is None:
            borders = OxmlElement('w:tblBorders')
            props.append(borders)
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            element = OxmlElement('w:' + edge)
            for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                element.set(qn('w:' + name), value)
            borders.append(element)
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E8EDF1')
                    cell._tc.get_or_add_tcPr().append(shade)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
    builder.native.stable_docx(doc, target)


if __name__ == '__main__':
    builder.word = word
    builder.CASES = CASES
    builder.ATTACHMENTS = {case['slug']: case.get('attachments', {}) for case in CASES}
    # Die fertig redigierte README gehört der Repository-Integration. Ein
    # reproduzierbarer Neubau der Originale soll sie nicht überschreiben.
    builder.readme = lambda case, directory: None
    builder.main()
