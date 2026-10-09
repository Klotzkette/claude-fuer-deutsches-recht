#!/usr/bin/env python3
"""Prüft Vertragsakte, Kapitalrechnung und unveränderte Downloadpakete."""

import argparse
import csv
from email import policy
from email.parser import BytesParser
import hashlib
import importlib.util
import io
import json
import re
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from testakte_disclaimer import NOTICE_BYTES
from testakte_file_filter import _safe_text, include_in_working_dump
from release_routing import (
    RELEASE_BASE, case_asset_url, companion_case_groups, companion_cases,
    companion_tag, marketplace_version, scoped_asset_url,
)

ROOT = Path(__file__).resolve().parents[1]
NAME = "gesellschaftervereinbarung"
SLUG = "gesellschaftervereinbarung-drohnenfriseur-berlin"
CASE = ROOT / "testakten" / SLUG
PLUGIN = ROOT / NAME
DIST = None


def originals():
    return sorted(p for p in CASE.rglob("*") if include_in_working_dump(p, CASE))


class AgreementTests(unittest.TestCase):
    def test_unpublished_case_follows_partition_and_component_pin_stays_stable(self):
        slug = "weltraumrecht-satellitenschwarm-startplatz-kueste"
        part = next(
            number for number, group in enumerate(companion_case_groups(companion_cases()), 1)
            if slug in group
        )
        for suffix in ("", "-einzelpdfs"):
            asset = f"testakte-{slug}{suffix}.zip"
            self.assertIsNone(scoped_asset_url(asset))
            self.assertEqual(
                f"{RELEASE_BASE}/download/{companion_tag(marketplace_version(), part)}/{asset}",
                case_asset_url(slug, suffix),
            )
            self.assertIn(
                "/gesellschaftervereinbarung-v1.1.0/", case_asset_url(SLUG, suffix)
            )

    def test_manifests_and_skills(self):
        manifests = [
            json.loads((PLUGIN / p).read_text())
            for p in (
                "plugin.json",
                ".claude-plugin/plugin.json",
                ".codex-plugin/plugin.json",
            )
        ]
        self.assertEqual(
            {(m["name"], m["version"], m["description"]) for m in manifests},
            {(NAME, "1.1.0", manifests[0]["description"])},
        )
        skills = list((PLUGIN / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 11)
        for skill in skills:
            text = skill.read_text()
            for heading in range(1, 7):
                self.assertRegex(text, rf"(?m)^#{{1,2}} {heading} ", str(skill))
            self.assertIn("Paragraf", text, str(skill))
            self.assertIn("BGH", text, str(skill))
        for kind in ("schnellstart", "hauptproblem"):
            data = (PLUGIN / f"{NAME}-{kind}.md").read_bytes()
            self.assertLessEqual(len(data), 7500)
            self.assertLessEqual(len(data.decode()), 7500)

    def test_complete_native_file_set(self):
        files = originals()
        self.assertEqual(len(files), 46)
        self.assertEqual(
            {
                ext: sum(p.suffix == ext for p in files)
                for ext in (".docx", ".pdf", ".eml", ".csv", ".txt", ".xlsx")
            },
            {".docx": 2, ".pdf": 23, ".eml": 12, ".csv": 5, ".txt": 3, ".xlsx": 1},
        )
        for source in CASE.glob("*.csv"):
            rows = list(
                csv.reader(
                    io.StringIO(source.read_text(encoding="utf-8")), delimiter=";"
                )
            )
            self.assertGreaterEqual(len(rows), 5)
            self.assertEqual(len({len(row) for row in rows}), 1)

    def test_template_exception_remains_narrow(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp) / SLUG
            directory.mkdir()
            approved = directory / "02_Gesellschaftervereinbarung_Ausfuellfassung.docx"
            other = directory / "03_Entwurf.docx"
            document = Document()
            document.add_paragraph("Vereinbart wird der Vollzug am [Datum].")
            document.save(approved)
            document.save(other)
            self.assertTrue(include_in_working_dump(approved, directory))
            self.assertFalse(include_in_working_dump(other, directory))
            document.add_paragraph("Musterlösung für den Prüfer")
            document.save(approved)
            _safe_text.cache_clear()
            self.assertFalse(include_in_working_dump(approved, directory))

    def test_mail_headers_and_binary_attachments(self):
        attachments = 0
        for path in CASE.glob("*.eml"):
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            for header in (
                "From",
                "To",
                "Date",
                "Subject",
                "Message-ID",
                "MIME-Version",
            ):
                self.assertTrue(message[header], f"{path.name}: {header}")
            self.assertFalse(message.defects, path.name)
            for part in message.iter_attachments():
                attachments += 1
                self.assertEqual(
                    part.get_payload(decode=True),
                    (CASE / part.get_filename()).read_bytes(),
                    path.name,
                )
        self.assertEqual(attachments, 12)

    def test_capital_and_cached_formulas(self):
        capital_skill = (
            PLUGIN / "skills/kapital-und-tranchen-abgleichen/SKILL.md"
        ).read_text()
        self.assertIn(
            "Ausgabeaufgeld fällt unter Paragraf 272 Absatz 2 Nummer 1 HGB",
            capital_skill,
        )
        source = CASE / "22_Kapital_und_Finanzierungsplan.xlsx"
        wb = load_workbook(source, data_only=True)
        self.assertEqual(wb.sheetnames, ["Kapital", "Tranchen", "Budget", "Historie", "Darlehen", "Wandlung", "Kontojournal"])
        self.assertEqual(
            [wb["Kapital"][cell].value for cell in ("C13", "D13", "E13", "F13", "G13")],
            [50000, 12000, 62000, 8000, 70000],
        )
        for cell in ("H13", "I13"):
            self.assertAlmostEqual(wb["Kapital"][cell].value, 1)
        self.assertEqual(wb["Tranchen"]["D14"].value, 10000000)
        self.assertEqual(wb["Tranchen"]["H14"].value, 9980000)
        self.assertEqual(wb["Budget"]["D13"].value, 10000000)
        self.assertEqual([wb["Historie"][f"I{r}"].value for r in range(6, 11)], [1000, 25000, 30000, 50000, 50000])
        self.assertEqual(wb["Darlehen"]["C12"].value, 470000)
        days = (date(2026, 10, 1) - date(2024, 5, 1)).days
        interest = (Decimal(120000) * Decimal("0.06") * days / 365).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        self.assertEqual(wb["Darlehen"]["F7"].value, days)
        self.assertEqual(Decimal(str(wb["Darlehen"]["G7"].value)), interest)
        self.assertEqual(wb["Darlehen"]["G12"].value, 27630.4)
        self.assertEqual(wb["Wandlung"]["C12"].value, 343)
        self.assertEqual(wb["Wandlung"]["C13"].value, 218.08)
        self.assertEqual(wb["Wandlung"]["C15"].value, 70343)
        self.assertEqual(wb["Kontojournal"]["F17"].value, 80000)
        self.assertTrue(all(wb["Kontojournal"][f"H{r}"].value == 0 for r in range(6,18)))
        formulas = load_workbook(source, data_only=False)
        for sheet in formulas:
            self.assertTrue(sheet.print_area)
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "e")
                    if cell.data_type == "f":
                        cached = wb[sheet.title][cell.coordinate]
                        self.assertIsNotNone(cached.value, cell.coordinate)
                        self.assertNotEqual(cached.data_type, "e")
        self.assertEqual(formulas["Tranchen"]["H7"].value, "=D7-G7")
        wb.close()
        formulas.close()

    def test_complete_pdf_and_editable_documents(self):
        combined = PdfReader(CASE / "gesamt-pdf" / f"{SLUG}_gesamt.pdf")
        self.assertGreaterEqual(len(combined.pages), 100)
        text = "\n".join(page.extract_text() or "" for page in combined.pages)
        for phrase in (
            "32.3 Vollständige Abschlussfelder",
            "Kunigunde",
            "Finanzierungsplan",
        ):
            self.assertTrue(phrase in text, f"PDF-Inhalt fehlt: {phrase}")
        self.assertNotIn("Diese Testakte wurde mit KI generiert", text)
        for path in CASE.glob("*.docx"):
            document = Document(path)
            self.assertGreater(sum(len(p.text) for p in document.paragraphs), 10000)
            self.assertEqual(document.styles["Normal"].font.name, "Times New Roman")
            self.assertEqual(document.styles["Normal"].font.size.pt, 11)

    def test_fillable_contract_structure(self):
        document = Document(CASE / "02_Gesellschaftervereinbarung_Ausfuellfassung.docx")
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertGreater(len(text), 75000)
        self.assertGreater(len(set(re.findall(r"\[[^\]]+\]", text))), 200)
        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
        for heading in headings:
            self.assertRegex(heading, r"^\d+(?:\.\d+)* ")
        self.assertEqual([int(p.text.split()[0]) for p in document.paragraphs if p.style.name == "Heading 1"], list(range(1,33)))
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Header", "Footer"):
            self.assertEqual(document.styles[name].font.name, "Times New Roman")
            self.assertEqual(document.styles[name].font.size.pt, 11)
        self.assertIn("keine Zahlung auf das Stammkapital", text)
        self.assertIn("Rechte außenstehender Gläubiger", text)
        signature = next(p for p in document.paragraphs if p.text == "32.3 Vollständige Abschlussfelder")
        self.assertTrue(signature.paragraph_format.page_break_before)
        journal = PdfReader(CASE / "35_Zahlungsjournal_Februar_Maerz.pdf")
        self.assertEqual(len(journal.pages), 1)

    def test_historical_cash_and_capital_are_separate(self):
        with (CASE / "19_Historische_Einzahlungen.csv").open(encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle, delimiter=";"))
        for person in ("Kunigunde Wolkenberger", "Kilian Funkenschlag"):
            entries = [r for r in rows if r["Beteiligter"] == person]
            self.assertEqual([(r["Datum"], r["Betrag_EUR"]) for r in entries], [("2023-06-12", "500"), ("2023-09-20", "12000")])
        with (CASE / "44_Bankjournal_Fruehjahr.csv").open(encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle, delimiter=";"))
        balance = Decimal(rows[0]["Saldo_EUR"])
        for row in rows[1:]:
            balance += Decimal(row["Zugang_EUR"]) - Decimal(row["Abgang_EUR"])
            self.assertEqual(balance, Decimal(row["Saldo_EUR"]))
        self.assertEqual(balance, 80000)

    def test_release_packages(self):
        if DIST is None:
            self.skipTest("Paketprüfung nur mit --dist")
        expected = {p.name for p in originals()}
        with zipfile.ZipFile(DIST / f"testakte-{SLUG}.zip") as archive:
            self.assertEqual(
                set(archive.namelist()), expected | {"README.txt", f"{SLUG}_gesamt.pdf"}
            )
            self.assertTrue(archive.read("README.txt").startswith(NOTICE_BYTES))
            self.assertEqual(
                archive.read(f"{SLUG}_gesamt.pdf"),
                (CASE / "gesamt-pdf" / f"{SLUG}_gesamt.pdf").read_bytes(),
            )
            for source in originals():
                self.assertEqual(archive.read(source.name), source.read_bytes())
        with zipfile.ZipFile(DIST / f"testakte-{SLUG}-einzelpdfs.zip") as archive:
            self.assertEqual(len(archive.namelist()), 47)
            self.assertTrue(archive.read("README.txt").startswith(NOTICE_BYTES))
            for name in archive.namelist():
                self.assertNotIn("/", name)
                if name != "README.txt":
                    self.assertTrue(name.endswith(".pdf"))
                    self.assertGreater(
                        len(PdfReader(io.BytesIO(archive.read(name))).pages), 0
                    )
        for portable in (False, True):
            prefix = f"{NAME}/" if portable else ""
            with zipfile.ZipFile(
                DIST / f"{NAME}{'-portable' if portable else ''}.zip"
            ) as archive:
                names = archive.namelist()
                self.assertEqual(sum(n.endswith("/SKILL.md") for n in names), 11)
                self.assertIn(prefix + "LICENSE", names)
                for kind in ("werkstatt", "schnellstart", "hauptproblem"):
                    self.assertNotIn(prefix + f"{NAME}-{kind}.md", names)
                for source in (PLUGIN / "skills").glob("*/SKILL.md"):
                    self.assertEqual(
                        archive.read(prefix + source.relative_to(PLUGIN).as_posix()),
                        source.read_bytes(),
                    )
        checksums = (DIST / "checksums-sha256.txt").read_text().splitlines()
        self.assertEqual(len(checksums), 10)
        for line in checksums:
            digest, name = line.split("  ", 1)
            self.assertEqual(
                hashlib.sha256((DIST / name).read_bytes()).hexdigest(), digest
            )
        # Die zentralen Paketprüfer ohne die nur beim Gesamtrelease nötigen Sammelarchive.
        for filename, suffix in (
            ("validate-testakten-release-zips.py", ""),
            ("validate-testakten-einzelpdf-zips.py", "-einzelpdfs"),
        ):
            spec = importlib.util.spec_from_file_location(
                filename, ROOT / "scripts" / filename
            )
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            path = DIST / f"testakte-{SLUG}{suffix}.zip"
            if suffix:
                actual = module.zip_entries(path, expected_suffix=".pdf")
                expected_names = ["README.txt", *module.expected_arcnames(CASE)]
            else:
                actual = module.zip_entries(path, require_notice=True)
                expected_names = module.expected_entries(CASE)
            module.assert_same(path.name, expected_names, actual)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dist", type=Path)
    DIST = parser.parse_args().dist
    unittest.main(argv=[__file__], verbosity=2)
