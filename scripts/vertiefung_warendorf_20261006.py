#!/usr/bin/env python3
"""Additive Vertiefung WH26; ausschließlich Warendorf, Originale 01–30 unverändert."""
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
import hashlib, json, re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.sax.saxutils import escape
from akten_build_runtime import serif_font_path
from akten_docx_format import separate_section_headings

ROOT=Path(__file__).resolve().parents[1]
CASE=ROOT/'testakten/bauwirtschaft-baumanagement-werkhalle-warendorf'
QA=Path('/tmp/bauwirtschaft-vertiefung-20261006/warendorf')
QA.mkdir(parents=True,exist_ok=True)
stamp=datetime(2026,9,25,16)
pdfmetrics.registerFont(TTFont('TNR',str(serif_font_path())))
pdfmetrics.registerFont(TTFont('TNRB',str(serif_font_path(bold=True))))
before={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in CASE.iterdir() if p.is_file() and re.match(r'\d\d_',p.name) and int(p.name[:2])<=30}

def doc(name,issuer,receiver,date,title,blocks):
    d=Document(); s=d.sections[0]
    s.page_width,s.page_height=Cm(21),Cm(29.7)
    s.top_margin,s.bottom_margin,s.left_margin,s.right_margin=Cm(1.9),Cm(1.9),Cm(2.2),Cm(2.2)
    for key in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[key]; st.font.name='Times New Roman'; st.font.size=Pt(11); st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().rFonts
        for a in list(fonts.attrib):
            if a.endswith('Theme'): del fonts.attrib[a]
        for a in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+a),'Times New Roman')
        st.paragraph_format.space_after=Pt(7); st.paragraph_format.line_spacing=1.04
    d.styles['Title'].font.size=Pt(15); d.styles['Title'].font.bold=True
    d.styles['Heading 1'].font.bold=True
    for b in d.styles.element.xpath('.//w:pBdr'): b.getparent().remove(b)
    for t in [issuer,receiver,date]: d.add_paragraph(t)
    d.add_paragraph(title,'Title')
    for h,text in blocks:
        d.add_paragraph(h,'Heading 1')
        for p in text.split('\n\n'): d.add_paragraph(p)
    d.core_properties.author=d.core_properties.last_modified_by='Klotzkette'
    d.core_properties.title=title; d.core_properties.language='de-DE'; d.core_properties.created=d.core_properties.modified=stamp
    separate_section_headings(d)
    d.save(CASE/name)

def pdf(name,issuer,receiver,date,title,blocks):
    n=ParagraphStyle('N',fontName='TNR',fontSize=11,leading=14,spaceAfter=8)
    h=ParagraphStyle('H',parent=n,fontName='TNRB',spaceBefore=5,spaceAfter=8)
    t=ParagraphStyle('T',parent=h,fontSize=15,leading=18,spaceAfter=13)
    flow=[Paragraph(escape(x).replace('\n','<br/>'),n) for x in [issuer,receiver,date]]+[Spacer(1,6),Paragraph(escape(title),t)]
    for heading,text in blocks:
        paras=text.split('\n\n')
        flow.append(KeepTogether([Paragraph(escape(heading),h),Paragraph(escape(paras[0]),n)]))
        flow.extend(Paragraph(escape(p),n) for p in paras[1:])
    def foot(c,d):
        c.setFont('TNR',8); c.drawString(62,30,f'WH26 | {name[:2]} | Seite {d.page}')
    SimpleDocTemplate(str(CASE/name),pagesize=A4,leftMargin=62,rightMargin=62,topMargin=45,bottomMargin=45,title=title,author='Klotzkette').build(flow,onFirstPage=foot,onLaterPages=foot)

def mail(name,sender,recipient,date,subject,body,attachment=None):
    m=EmailMessage(policy=SMTP)
    m['From']=sender; m['To']=recipient; m['Date']=format_datetime(datetime.fromisoformat(date+'+02:00'))
    address=sender.split('<')[1].split('>')[0]
    m['Return-Path']='<'+address+'>'
    m['Received']=f'from mail.{address.split("@")[1]} by archive.hagedorn.example with ESMTPS id WH26-{name[:2]}; '+m['Date']
    m['Subject']=subject; m['Message-ID']=f'<wh26-{name[:2]}-20260925@hagedorn.example>'
    m['Content-Language']='de-DE'; m.set_content(body.strip()+'\n')
    if attachment:
        p=CASE/attachment; m.add_attachment(p.read_bytes(),maintype='application',subtype='pdf',filename=p.name)
    (CASE/name).write_bytes(m.as_bytes())

mail('31_Hersteller_Schutzbaugruppe.eml','Petra Drees <disposition@schaltwerk-ostwestfalen.example>','Nils Breden <nils@elektro-breden.example>','2026-09-24T14:42:00','SW 26841 / Schutzbaugruppe und Endprüfung der Station WH26', '''Guten Tag Herr Breden,

für die Station SW 26841 fehlt weiterhin das Schutzmodul SPM-630, Serienreservierung 26-4418. Die ursprünglich für den 21. September bestätigte Wareneingangsprüfung konnte nicht stattfinden. Unser Unterlieferant nennt nun den 2. November für die Anlieferung. Diese Mitteilung ist eine Dispositionsauskunft. Einen Versandnachweis haben wir noch nicht erhalten.

Der Transformator und das Stationsgehäuse stehen seit 18. September in unserer Halle 3. Die Vorprüfung der Verdrahtung ist abgeschlossen; die komplette Funktionsprüfung mit dem Schutzmodul steht aus. Unser Prüfstand ist für den 5. November reserviert. Bei erfolgreicher Prüfung kann die Station am 6. November verladen werden. Ihre Anlieferung beim Kunden am 9. November ist damit darstellbar. Ein früheres Verladen ohne das Modul verkürzt die offene Funktionsprüfung nicht.

Wir haben ein Ersatzmodul aus einer anderen Baureihe untersucht. Dieses wäre mechanisch einbaubar, erfordert aber eine erneute Parametrierung und Freigabe durch die Elektroplanung. Die Herstellerfreigabe hierfür liegt nicht vor. Bitte geben Sie die Angabe „Ersatz vorhanden“ aus dem Telefonat mit unserem Lager vom 23. September deshalb nicht als bestätigte Alternative weiter.

Bitte bestätigen Sie bis 28. September, ob die ursprüngliche Schutzparametrierung unverändert bleibt. Änderungen des Anschlusskonzepts würden eine neue Prüfung des reservierten Zeitfensters auslösen. Die Entscheidung über die Lieferbedingungen Ihres Vertrags mit der Bauherrin treffen wir mit dieser Nachricht nicht. Ich melde den tatsächlichen Eingang des Moduls gesondert.

Freundliche Grüße
Petra Drees
Disposition Schaltwerk Ostwestfalen GmbH
Weiterleitung an Projektcontrolling Hagedorn am 25.09.2026 um 07:56 Uhr durch Nils Breden.''')

doc('32_Rueckfrage_Lieferkette.docx','Hagedorn Präzisionsteile GmbH\nEinkauf Leon Rensing','Elektro Breden GmbH\nNils Breden','25. September 2026, 09:45 Uhr','Lieferkette und Zwischenversorgung',[
('1 Rückmeldung zum Liefertermin','Sehr geehrter Herr Breden, wir haben Ihre Lieferfortschreibung vom 24. September und die weitergeleitete Nachricht von Frau Drees erhalten. Bitte benennen Sie uns bis 28. September, 09:00 Uhr, welche Schritte zwischen Eingang des Schutzmoduls am 2. November und der Anlieferung am 9. November noch offen sind. Wir benötigen außerdem die Rückmeldung des Unterlieferanten und den nachgereichten Versandnachweis, sobald dieser vorliegt. Die Herstellerinformation nennt eine Reservierung des Prüfstands, jedoch noch keinen abgeschlossenen Test.'),
('2 Prüfung der Alternativen','Bitte legen Sie offen, welchen Einfluss eine geänderte Schutzparametrierung auf die Prüfung und den Netztermin hätte. Unsere Geschäftsführung möchte eine frühere Teilanlieferung nur weiterverfolgen, wenn Sie den tatsächlichen Zeitgewinn und die zusätzlichen Anschlussarbeiten beziffern können. Die Nutzung eines anderen Moduls ist bislang weder technisch geprüft noch beauftragt.\n\nFür das mobile Angebot MB 0925 benötigen wir die Darstellung der Versorgung einer einzelnen Linie einschließlich Hallenverbrauchern. Die bloße Addition der Nennwerte genügt uns für den Startversuch nicht. Bitte nennen Sie insbesondere den anzusetzenden Anlaufstrom, die Umschaltfolge und die Person, die das Prüfprotokoll am Aufstelltag unterschreiben wird.'),
('3 Kosten und Dokumentation','Bitte trennen Sie den vorhandenen Einschichtpreis von Mehrkosten für Arbeit nach 20:00 Uhr und von einer Nutzung nach dem 22. November. In dem Betrag von 36.000,00 EUR sind sechs Wochen und der bisher angebotene Betriebsstoffansatz bereits enthalten. Wir möchten dieselben Mengen nicht in einem Ergänzungsangebot noch einmal vergüten. Die laufende Mietfläche wird über eine andere Kostenstelle geführt.\n\nWir halten die zugesagten ursprünglichen Daten aus der Bestellung EB 260318 in unserer Ablage fest. Mit der Anforderung weiterer Angaben stimmen wir keiner Vertragsänderung und keiner zusätzlichen Vergütung zu. Frau Hagedorn entscheidet am 28. September nach Sichtung der Unterlagen. Bitte bestätigen Sie den Eingang dieses Schreibens.\n\nMit freundlichen Grüßen\nLeon Rensing\nEinkauf')])

pdf('33_Anschlusspruefung_Vorbehalte.pdf','Elektro Breden GmbH\nElektroplanung Nils Breden','Hagedorn Präzisionsteile GmbH\nJana Feldkamp','25. September 2026, 12:10 Uhr','Anschlussprüfung der mobilen Versorgung',[
('1 Gegenstand der Vorprüfung','Wir haben die Geräteblätter vom 17. September und die Anschlusswerte aus unserem Blatt vom 22. September am Schreibtisch abgeglichen. Dies ist keine Messung an einer aufgebauten Anlage. Linie 1 ist mit 220 kVA und die allgemeinen Hallenverbraucher mit 20 kVA angesetzt. Gegenüber der angebotenen mobilen Versorgung mit 250 kVA verbleiben rechnerisch 10 kVA. Linie 2 mit weiteren 230 kVA kann dabei nicht zugeschaltet werden.'),
('2 Noch fehlende technische Unterlagen','Der Maschinenbauer hat für den Hauptspindelanlauf bislang nur den Nennwert mitgeteilt. Die Stromkurve des Starts wird bis 28. September angefordert. Wir können deshalb weder den Spannungseinbruch noch die Auslösung beim Anlauf abschließend bewerten. Bis zur Klärung bleibt der Betrieb der CNC-Linie offen; die rechnerischen 10 kVA sind keine zugesagte Anlaufreserve.\n\nDas Schaltbild MP-WH26-02 wird bis 28. September, 08:30 Uhr, fertiggestellt. Es soll eine mechanisch verriegelte Umschaltung und getrennte Einspeisepunkte vorsehen. Die Nachricht des Anschlussservices vom 25. September fordert zusätzlich den Schutzprüfplan. Beide Unterlagen sind noch nicht freigegeben. Der Anschlussservice soll danach den Anschluss am 12. Oktober prüfen.'),
('3 Ablauf am Aufstelltag','Nach Aufstellung sind die Schutzprüfung, die Prüfung der Verriegelung und ein belasteter Startversuch mit Linie 1 vorgesehen. Elektro Breden stellt die Elektrofachkraft, der Maschinenbauer führt den Startversuch. Wir bitten Hagedorn, bis zum 28. September einen Ansprechpartner für den gesperrten Maschinenbereich zu benennen. Die Verbindung mit dem Netz darf erst nach dokumentierter Umschaltfreigabe hergestellt werden. Dies ist unsere Vorgabe für das angebotene Anschlusskonzept.'),
('4 Bezug zu Kosten und Termin','Die bezeichneten Anschlussprüfungen gehören zum Grundangebot MB 0925. Wiederholte Versuche aufgrund geänderter Maschinenparameter würden wir vor Ausführung gesondert anbieten. Ein Datum für die Produktionsfreigabe bestätigen wir heute nicht. Die beiden Zeitansätze aus der Baufolge, 15 Kalendertage Versuche und sechs Kalendertage Einweisung, bleiben Planannahmen.\n\nNils Breden\nElektroplanung')])

pdf('34_Ergaenzungsangebot_Spaetschicht.pdf','Elektro Breden GmbH\nHülsener Straße 31, 59302 Oelde','Hagedorn Präzisionsteile GmbH\nEinkauf Leon Rensing','25. September 2026, 13:05 Uhr','Ergänzungsangebot zur mobilen Versorgung',[
('1 Vorgeschlagener Umfang','Auf Ihre telefonische Anfrage bieten wir ergänzend zu MB 0925 zehn zusätzliche Abendfenster von jeweils drei Stunden zwischen 20:00 und 23:00 Uhr an. Die Tage würden im Oktober und November nach Abruf abgestimmt. Unsere Preisannahme setzt zwei Elektrofachkräfte je Fenster voraus. Die Ausführung hängt von einer gesonderten Beauftragung sowie den technischen und betrieblichen Freigaben ab. Eine behördliche Klärung der angegebenen Uhrzeiten liegt uns nicht vor.'),
('2 Zusätzliche Preise','Die Fachkräfte berechnen wir mit 78,00 EUR netto je Personenstunde. Für zehn Fenster, je drei Stunden und zwei Personen ergeben sich 4.680,00 EUR. Für die zusätzliche Sicherungswache rechnen wir zehnmal drei Stunden zu 42,00 EUR, mithin 1.260,00 EUR. Zusätzlichen Betriebsstoff für diese Fenster setzen wir mit 18,00 EUR je Betriebsstunde an, zusammen 540,00 EUR. Die zusätzliche Schallschutzaufstellung kostet einmalig 1.250,00 EUR. Das Ergänzungsangebot beträgt damit 7.730,00 EUR netto, zuzüglich 1.468,70 EUR Umsatzsteuer, insgesamt 9.198,70 EUR.\n\nDie Personensätze umfassen die von uns kalkulierten Zuschläge. Es sind keine Zuschläge nochmals auf diese Sätze aufzuschlagen. Die Betriebsstoffpauschale von 3.000,00 EUR im Grundangebot bleibt für den Einschichtbetrieb bestehen; die 540,00 EUR betreffen ausschließlich die zusätzlichen Abendstunden. Bei weniger abgerufenen Stunden rechnen wir die tatsächlichen Stunden und die einmalig ausgeführte Schallschutzaufstellung ab.'),
('3 Laufzeit und Abrechnung','Der bisherige Mietzeitraum endet am 22. November. Eine Verlängerung kostet je zusätzlich angefangener Woche 4.000,00 EUR Miete und 500,00 EUR Betriebsstoff für Einschichtbetrieb. Wir halten höchstens zwei zusätzliche Wochen disponibel. Eine Verlängerung ist separat bis 13. November zu bestellen. Zusätzliche Abendstunden während der Verlängerung sind damit nicht abgedeckt.\n\nAuf das Ergänzungsangebot verlangen wir keine Vorauszahlung. Die ausgeführten Zusatzstunden und die Schallschutzaufstellung werden nach Rückbau mit unterzeichneten Tagesnachweisen abgerechnet. Die Vorauszahlung von 30 Prozent aus MB 0925 bezieht sich ausschließlich auf den Grundpreis von 36.000,00 EUR. Unsere Preise bleiben bis 28. September, 14:00 Uhr, gebunden.\n\nNils Breden\nGeschäftsführer')])

mail('35_Anfrage_Abendbetrieb.eml','Jana Feldkamp <projekt@hagedorn.example>','Umweltstelle <umwelt@warendorf.example>','2026-09-25T13:22:00','WH26 / Klärung von Versuchsbetrieb zwischen 20 und 23 Uhr', '''Sehr geehrte Damen und Herren,

wir bitten um Auskunft zum vorgesehenen Versuchsbetrieb mit mobiler Stromversorgung auf unserem Grundstück Werkstraße 18. Elektro Breden hat zehn zusätzliche Zeitfenster von jeweils 20:00 bis 23:00 Uhr angeboten. Festgelegt sind noch keine einzelnen Tage. Die Fenster sollen im Oktober und November 2026 liegen. Eine Bestellung haben wir bisher nicht erteilt.

Der technische Zweck ist die Erprobung einer CNC-Linie in der neuen Halle. Das Aggregat soll neben der östlichen Hallenwand stehen. Die nächstgelegene Wohnnutzung liegt nach unserem Lageplan etwa 85 Meter entfernt. Uns fehlen bislang belastbare Angaben zum Schallpegel am Immissionsort, zur Dauer der Lastwechsel und zu den Rangierbewegungen bei der Kraftstoffversorgung. Der Lieferant hat eine zusätzliche Schallschutzaufstellung angeboten, hierzu aber noch kein Mess- oder Berechnungsblatt übersandt.

Bitte teilen Sie uns mit, welche Stelle den konkreten Vorgang bearbeitet und welche Unterlagen für eine Prüfung der vorgesehenen Zeitfenster erforderlich sind. Bitte nennen Sie auch, ob zwischen Aufstellung, baulichen Anschlussarbeiten und anschließendem Versuchsbetrieb zu unterscheiden ist. Wir möchten Ihre Rückmeldung vor einer zeitlichen Zusage gegenüber den ausführenden Firmen berücksichtigen.

Elektro Breden stellt den Schutzprüfplan und das Schaltbild bis Montag bereit. Die noch offene elektrische Anschlussfreigabe behandeln wir getrennt von Ihrer Rückmeldung. Eine Nachricht über den Eingang der Anfrage würde uns bei der Terminabstimmung helfen.

Mit freundlichen Grüßen
Jana Feldkamp
Projektcontrolling Hagedorn Präzisionsteile GmbH

Ablagevermerk 25.09.2026, 15:45 Uhr: Versand im Projektpostfach dokumentiert. Eine Antwort liegt bis zum Stichtag nicht vor.''')

doc('36_Freigabevermerk_Vorbereitung.docx','Hagedorn Präzisionsteile GmbH\nMaren Hagedorn Geschäftsführerin','Jana Feldkamp und Leon Rensing\nProjekt WH26','25. September 2026, 14:10 Uhr','Vorbereitung der Energieentscheidung',[
('1 Erlaubte Vorbereitung','Bitte führen Sie die schriftlichen Rückfragen zur Lieferkette, zum Netzanschluss und zu den angebotenen Abendfenstern bis Montag zusammen. Sie dürfen Elektro Breden um eine Verlängerung der Angebotsbindung bitten. Eine kostenpflichtige Reservierung dürfen Sie damit nicht auslösen. Die technische Planung und die Anschlussprüfung aus dem vorhandenen Angebot sind zunächst preislich zu bestätigen.'),
('2 Entscheidung über Aufträge','Ich habe MB 0925 und das Ergänzungsangebot vom heutigen Tag erhalten. Keines der beiden Angebote ist angenommen. Bitte legen Sie mir die Grundkosten von 36.000,00 EUR, die gesonderten Zusatzstunden und jede verlängerte Mietwoche jeweils erkennbar vor. Die vorbereitenden Rückfragen sind keine Freigabe der 30-prozentigen Vorauszahlung.\n\nFür die Kranforderung KM N03 bleibt es bei der offenen Klärung. Herr Rensing soll uns die Grundlage der behaupteten durchgehenden Bereitstellung erläutern lassen. Die von der Bauleitung unterzeichneten Einsätze genügen mir für eine Entscheidung über die gesamte Forderung noch nicht. Der Budgetbeschluss vom Januar bleibt unverändert.'),
('3 Zahlungsdisposition und Restleistungen','Frau Voß soll den offenen Westkamp-Abschlag in der Zahlungsvorschau wahlweise mit der vorsorglichen Septemberreserve und dem Zahlungsdatum 8. Oktober ausweisen. Beide Darstellungen müssen denselben Endbestand ergeben, solange alle übrigen Ansätze gleich bleiben. Eine Bankausführung bestätige ich mit diesem Vermerk nicht.\n\nBitte prüfen Sie ferner das neue Angebot zur Druckluftverteilung gegen den bisherigen Restkostenansatz von 30.000,00 EUR. Eine Bestellung der gesamten angebotenen Summe zusätzlich zu diesem Ansatz würde denselben Leistungsumfang doppelt erfassen. Der Einkauf soll die angebotenen Mengen und die Abgrenzung zu Elektro Breden klären.'),
('4 Nächste Vorlage','Ich erwarte am 28. September um 10:00 Uhr eine Gegenüberstellung der noch offenen Voraussetzungen, der bedingten Termine und der tatsächlichen Zahlungsanlässe. Die Antwort auf die Anfrage zum Abendbetrieb kann später eintreffen. In diesem Fall soll die Vorlage diese offene Voraussetzung ausdrücklich benennen.\n\nMaren Hagedorn\nGeschäftsführerin')])

doc('37_Rechnungseingang_Zahlungslauf.docx','Hagedorn Präzisionsteile GmbH\nFinanzen Elena Voß','Maren Hagedorn und Jana Feldkamp\nProjektkonto WH26','25. September 2026, 14:35 Uhr','Rechnung WH 260924 im Zahlungslauf',[
('1 Rechnungseingang und Prüfung','Die Rechnung WH 260924 ging am 24. September um 10:18 Uhr im zentralen Rechnungsfach ein. Sie wurde am 25. September unter den Belegnummern WH-03-2 und WH-04-2 erfasst. Frau Feldkamp hat den Leistungsnachweis vom 23. September für den Septemberzuwachs von 160.000,00 EUR netto bestätigt. Der Zahlungsbetrag beträgt 190.400,00 EUR brutto. Eine Bankbuchung liegt hierzu nicht vor.'),
('2 Berechnung des Zahlungsdatums','Im Vertrag vom 19. Februar sind 14 Kalendertage nach Eingang der prüfbaren Rechnung vereinbart. Im Rechnungseingang ist der 24. September als vollständiger Eingang erfasst; der auf der Rechnung genannte 8. Oktober entspricht diesem Ansatz. Für die Disposition führen wir daher den 8. Oktober. Wird die Erfassung des Rechnungseingangs berichtigt, muss der Zahlungsansatz mitgeprüft werden.\n\nDas alte Monatsblatt behandelt denselben Betrag vorsorglich schon im September. Es bildet damit eine Liquiditätsreserve vor dem Zahlungstag. In einer nach Fälligkeit geführten Fortführung verschieben sich 190.400,00 EUR vom September in den Oktober. Der Betrag darf nicht in beiden Monaten abgezogen werden. Die zusätzliche Oktoberleistung von 260.000,00 EUR netto ist eine andere Planposition und bleibt bestehen.'),
('3 Bank und Freigabe','Der Kontostand vom heutigen Vormittag beträgt 512.400,00 EUR. Die vorgesehenen Abrufe von 400.000,00 EUR im Oktober und 600.000,00 EUR im November bleiben unverändert. Sie sind in der Vorschau geplante Zuflüsse. Die Bankansicht vom 25. September bleibt der Beleg für den aktuellen Bestand.\n\nIch habe den Betrag für einen Zahlungslauf am 8. Oktober vorgemerkt. Die zweite Zeichnung im Bankportal und die endgültige Ausführungsbestätigung stehen aus. Bitte geben Sie uns die Zahlungsfreigabe rechtzeitig vor dem Lauf. Ein als Zahlungsvorschlag angelegter Datensatz ist noch keine erfolgte Zahlung.\n\nElena Voß\nFinanzen')])

pdf('38_Angebot_Druckluftverteilung.pdf','Drucklufttechnik Reuter GmbH\nMontageplanung Thomas Reuter','Hagedorn Präzisionsteile GmbH\nLeon Rensing Einkauf','25. September 2026, 11:40 Uhr','Angebot DR 260925 zur Druckluftverteilung',[
('1 Umfang und Mengen','Wir bieten die werkstattseitige Druckluftverteilung ab Ausgang des vorhandenen Kompressors an. Grundlage ist die Begehung vom 23. September mit Herrn Rensing. Unser Angebot umfasst 160 Meter Aluminiumrohr einschließlich Halterung zu 96,00 EUR je Meter, zusammen 15.360,00 EUR. Zwölf Maschinenabgänge einschließlich Absperrung kosten je 620,00 EUR, zusammen 7.440,00 EUR. Für Montage und Dichtheitsprüfung setzen wir 80 Stunden zu 85,00 EUR, zusammen 6.800,00 EUR, an. Die Anschlussdokumentation berechnen wir pauschal mit 2.150,00 EUR.'),
('2 Preis und Abgrenzung','Der Gesamtpreis beträgt 31.750,00 EUR netto, zuzüglich 6.032,50 EUR Umsatzsteuer, insgesamt 37.782,50 EUR. Ein Kompressor, dessen elektrische Zuleitung sowie Fundamentarbeiten sind nicht enthalten. Der vorhandene Kompressor soll weitergenutzt werden. Die elektrische Zuleitung ist vor Montage mit Elektro Breden abzustimmen. Wir haben von dort noch keine bestätigte Schnittstellenzeichnung.\n\nHerr Rensing hat bei der Begehung einen bisherigen Ansatz von 30.000,00 EUR genannt. Unser Angebot ersetzt keine Bestellung und sagt nichts darüber aus, ob dieser Ansatz in Ihrer Kostenfortschreibung angepasst wird. Mengenänderungen nach einer Beauftragung werden vor Ausführung beschrieben und bepreist.'),
('3 Ablauf und Zahlung','Bei schriftlicher Bestellung bis 2. Oktober können wir ab 19. Oktober beginnen. Wir benötigen fünf Arbeitstage mit zugänglichem Trassenbereich. Eine gleichzeitige Montage über laufenden Maschinen bieten wir nicht an. Das ist mit der geplanten Maschinenaufstellung abzustimmen. Die Dokumentation übergeben wir nach der Dichtheitsprüfung.\n\nWir rechnen nach fertiggestellter Montage und Übergabe der Dokumentation ab; eine Vorauszahlung ist nicht vorgesehen. Der Rechnungsbetrag ist 14 Kalendertage nach vollständigem Rechnungseingang zahlbar. Das Angebot bleibt bis 2. Oktober gebunden. Eine Bestellung liegt uns am 25. September nicht vor.\n\nThomas Reuter\nMontageplanung')])

mail('39_Schnittstelle_Druckluft.eml','Nils Breden <nils@elektro-breden.example>','Leon Rensing <einkauf@hagedorn.example>','2026-09-25T15:05:00','WH26 / Grenze zwischen Technikauftrag und Druckluftangebot DR 260925', '''Guten Tag Herr Rensing,

die in unserem Technikauftrag enthaltene Zuleitung endet am Anschlusspunkt des vorhandenen Kompressors. Die Druckluftrohre und Maschinenabgänge aus DR 260925 sind in unserer Vergabesumme nicht enthalten. Ihr gesonderter Restkostenansatz für die Druckluftverteilung betrifft daher nach unserem Verständnis einen eigenen Leistungsumfang.

Bitte geben Sie dem Rohrleitungsbauer den vorgesehenen Maschinenstand aus Plan WH-M03 weiter. Die bislang besprochenen zwölf Abgänge passen zu dieser Fassung. Die Trassenhöhe oberhalb Linie 1 müssen beide Gewerke noch gemeinsam bestätigen. Ich kann heute nicht beurteilen, ob die angebotenen 160 Meter beim endgültigen Leitungsverlauf genügen. Dafür benötige ich die Trassenskizze von Herrn Reuter.

Für die Kostenrunde würde ich das neue Angebot zunächst neben den bisherigen Ansatz legen. Der neue Betrag von 31.750,00 EUR betrifft nach der Angebotsbeschreibung die gesamte Druckluftverteilung und nicht einen zusätzlichen Teil neben den 30.000,00 EUR. Die mengenmäßige Prüfung muss Ihr Einkauf vor einer Beauftragung abschließen.

Unsere Zuleitung ist im bestehenden Auftrag W05 und in den darin enthaltenen offenen Leistungen erfasst. Sie darf aus unserer Sicht nicht nochmals als Druckluftrestleistung angelegt werden. Die Trafostation bleibt ebenfalls Bestandteil dieser Vergabesumme; ihre offenen Abschlagsstufen sind keine zusätzliche Bestellung.

Freundliche Grüße
Nils Breden
Elektro Breden GmbH

Ablagevermerk von Leon Rensing um 15:25 Uhr: Trassenskizze bei Reuter angefordert. Angebot unverbindlich in der Einkaufsablage abgelegt; noch kein Bestellkennzeichen vergeben.''')

doc('40_Maschinenbauer_Versuchsfenster.docx','Maschinenservice Hellweg GmbH\nInbetriebnahmeplanung Sabine Krüger','Hagedorn Präzisionsteile GmbH\nJana Feldkamp','25. September 2026, 15:30 Uhr','Versuchsfenster für Linie 1 und Linie 2',[
('1 Bisheriger Ansatz','Wir nehmen Bezug auf den Ablaufvorschlag von Westkamp. Die dort genannten 15 Kalendertage für Versuche und sechs Kalendertage für Einweisung geben die im Termintelefonat vom 21. September verwendeten Ansätze zutreffend wieder. Beginn ist jeweils eine dokumentiert nutzbare Energieversorgung. Die Montage und Ausrichtung der Maschinen muss vorher abgeschlossen sein. Eine Reservierung eines Teams nur anhand eines nicht bestätigten Energiedatums ist noch keine Zusage für den Versuchsstart.'),
('2 Mobile Versorgung für Linie 1','Bei mobilem Betrieb kommt nur Linie 1 in Betracht. Die Anlaufkurve übermitteln wir Elektro Breden am 28. September. Bis deren Auswertung und der erfolgreichen Schutzprüfung können wir die Eignung der Anlage für den Start nicht bestätigen. Beide Linien gleichzeitig an dem angebotenen mobilen Aggregat zu testen, haben wir nicht angeboten.\n\nWir haben intern geprüft, ob zusätzliche Abendfenster die Versuchsphase auf zwölf Kalendertage verkürzen könnten. Dieser Wert ist eine Dispositionsannahme und steht unter dem Vorbehalt eines durchgehend verfügbaren Serviceteams, der nutzbaren Stromversorgung und der zulässigen Zeitfenster. Die sechs Kalendertage für Einweisung und Produktionsfreigabe bleiben in unserer Planung bestehen. Das Angebot von Elektro Breden betrifft dessen Personal und Versorgung; unsere zusätzlichen Servicekosten sind darin nicht enthalten. Einen Preis dafür können wir erst nach Bestätigung der einzelnen Tage nennen.'),
('3 Linie 2 und Abschaltung des Provisoriums','Linie 2 kann nach unserem jetzigen Stand erst nach der dauerhaften Versorgung in Betrieb genommen werden. Auch für sie setzen wir 15 Kalendertage Versuche und sechs Kalendertage Einweisung an. Die mobile Anlage kann deshalb nicht allein mit dem Ziel eines früheren Vollbetriebs bewertet werden. Der 16. November ist weiterhin ein bedingter Energieansatz.\n\nDer 22. November aus dem Mietangebot beendet die dort enthaltene Laufzeit. Bitte stimmen Sie mit uns einen sicheren Umschalttermin für Linie 1 ab, bevor Sie einen Rückbau beauftragen. Ob die mobile Anlage länger bleiben muss, hängt vom dokumentierten Anschluss und dem Umschaltversuch ab. Eine Verlängerung ist heute nicht bestätigt.'),
('4 Rückmeldung und Unterlagen','Bitte melden Sie am 28. September, welche Variante weiterverfolgt werden soll. Wir benötigen dann den bestätigten Energieplan und die einzelnen Zeitfenster. Die oben genannte mögliche Verkürzung darf bis zu unserer Teamzusage nur als Terminannahme geführt werden.\n\nSabine Krüger\nInbetriebnahmeplanung')])

after={n:hashlib.sha256((CASE/n).read_bytes()).hexdigest() for n in before}
assert before==after, 'Originale 01–30 verändert'
(QA/'originale_sha256.json').write_text(json.dumps(after,indent=2),encoding='utf-8')
readme=CASE/'README.md'
text=readme.read_text(encoding='utf-8')
text=text.replace('30 native Aktenstücke. Jede Datei bildet ein Dokument ab. Die Arbeitsmappen sind bearbeitbar; E-Mail-Anhänge entsprechen den separat enthaltenen Quelldateien.', '43 native Aktenstücke, darunter fünf bearbeitbare Excel-Arbeitsmappen. Die ursprünglichen 30 Unterlagen und ihre beiden Arbeitsmappen bleiben unverändert. Zehn zusätzliche Belege und drei neue Fortführungen vertiefen denselben Stichtag. Jede Datei bildet ein Dokument ab; E-Mail-Anhänge entsprechen den separat enthaltenen Quelldateien.')
names=['31_Hersteller_Schutzbaugruppe.eml','32_Rueckfrage_Lieferkette.docx','33_Anschlusspruefung_Vorbehalte.pdf','34_Ergaenzungsangebot_Spaetschicht.pdf','35_Anfrage_Abendbetrieb.eml','36_Freigabevermerk_Vorbereitung.docx','37_Rechnungseingang_Zahlungslauf.docx','38_Angebot_Druckluftverteilung.pdf','39_Schnittstelle_Druckluft.eml','40_Maschinenbauer_Versuchsfenster.docx','41_Energie_und_Freigaben.xlsx','42_Mobilversorgung_Kosten.xlsx','43_Zahlungen_und_Entscheidungen.xlsx']
if '| 31 |' not in text:
    text=text.replace('| 30 | `30_Finanzierung.eml` |','| 30 | `30_Finanzierung.eml` |\n'+'\n'.join(f'| {i} | `{n}` |' for i,n in enumerate(names,31)))
if '<!-- reserved-example-contacts -->' not in text:
    text=text.replace('> Diese Testakte wurde', '<!-- reserved-example-contacts -->\n\nDie ergänzten Kontaktadressen verwenden reservierte `.example`-Domains und sind nicht für einen Versand bestimmt.\n\n> Diese Testakte wurde',1)
if '## 1.6. Vertiefung und Workshop' not in text:
    text+='''

## 1.6. Vertiefung und Workshop

Die Ergänzung hält den Aktenstichtag 25. September 2026, 16:00 Uhr, ein. Sie führt keine späteren Freigaben, Zahlungen oder ausgeführten Versuche als Tatsachen ein. Der Hersteller erläutert die fehlende Schutzbaugruppe, Einkauf und Elektroplanung präzisieren die Alternativen, und der Maschinenbauer grenzt die mögliche Beschleunigung und seine noch offenen Zusatzkosten ab. Die Anfrage zum Abendbetrieb bleibt unbeantwortet.

| Arbeitsmappe | Weiterführung und Erklärung |
| --- | --- |
| 09 Ablauf und Zahlungen | Unveränderter ursprünglicher Plan ohne Provisorium und mit vorsorglicher Septemberreserve. Die ergänzende Datumslogik steht in 41 und 43. |
| 10 Kostenfortschreibung | Unveränderter Ausgangsstand für Nettoinvestition, Ist, Obligo und unbestellte Restleistungen. Die Angebotskalkulation 42 ändert keine Buchung. |
| 41 Energie und Freigaben | Zwei Blätter mit drei wählbaren Terminvarianten, getrenntem Anlauf beider Linien und sechs einzeln bestätigbaren Voraussetzungen. Eingaben, Kalendertagslogik und Quellen stehen neben den Werten. |
| 42 Mobilversorgung Kosten | Zwei Blätter mit Mengen, Personenstunden, Preisen, Mietwochen, Umsatzsteuer und Zahlungsstaffel. Der unbekannte Maschinenservice blockiert den Gesamtpreis nur in der gewählten Abendvariante. |
| 43 Zahlungen und Entscheidungen | Zwei Blätter mit Zahlungsdatum, Bruttorest, Einplanung und gesonderter Freigabe. Zwei vorbereitete neue Zeilen ermöglichen Fortführung; die Grenzen für weitere Zeilen sind erklärt. Fehlende Angaben und doppelte IDs bleiben sichtbar. |

Die folgenden Arbeitsaufträge verwenden die Akte als laufenden Vorgang. Sie geben keine Musterlösung vor.

1. Erstellen Sie aus 05 bis 07 und 31 bis 33 eine belastbare Lieferchronologie. Benennen Sie, welche Zusagen von welchem Unternehmen stammen und welche Nachweise noch fehlen.
2. Prüfen Sie die technische und kaufmännische Aussagekraft der mobilen Versorgung. Bearbeiten Sie in 41 jeweils nur eine neu belegte Voraussetzung und beobachten Sie, welche weiteren Schritte offen bleiben.
3. Kalkulieren Sie mit 34, 40 und 42 die zusätzlichen Abendfenster. Stellen Sie für die Geschäftsführung dar, welche Kosten aus den Angeboten bekannt sind und welche Angabe für eine vollständige Entscheidung noch benötigt wird.
4. Stimmen Sie Rechnung, Leistungsstand, Buchung und Bankbestand aus 21 bis 26 und 37 ab. Vergleichen Sie den frühen Reserveabzug mit dem Zahlungstag und dokumentieren Sie die zeitliche Verschiebung.
5. Prüfen Sie das Druckluftangebot 38 gegen Restkostenerhebung 13 und Schnittstellenmeldung 39. Führen Sie den Plan so fort, dass die bisherige Leistungsabgrenzung nachvollziehbar bleibt.
6. Entwerfen Sie die Vorlage für die Besprechung am 28. September mit konkreten Rückfragen, Zuständigkeiten und belegten Entscheidungsalternativen. Ein errechneter Termin oder Bankbestand allein ersetzt keine Freigabe.

## 1.7. Prüfung der Ergänzung

Die Angebotswerte übernehmen die Umsatzsteuerangaben der jeweiligen Belege. Die steuerliche Ausgangsannahme wurde am 6. Oktober 2026 erneut anhand der amtlichen Fassungen von [Paragraf 12 Absatz 1 UStG](https://www.gesetze-im-internet.de/ustg_1980/__12.html) und [Paragraf 13b Absatz 2 Nummer 4 sowie Absatz 5 UStG](https://www.gesetze-im-internet.de/ustg_1980/__13b.html) überprüft. Der Sachverhalt bleibt eine Bauherrin ohne eigene nachhaltige Bauleistungstätigkeit. Die Ergänzung bewertet weder die zivilrechtliche Verantwortlichkeit für den Lieferverzug noch die Zulässigkeit der angefragten Abendzeiten vorab.
'''
readme.write_text(text,encoding='utf-8')
print('10 neue Aktenstücke erstellt; 30 Originale unverändert.')
