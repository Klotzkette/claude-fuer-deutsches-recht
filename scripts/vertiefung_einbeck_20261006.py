#!/usr/bin/env python3
"""Fortschreibung der Einbecker Akte, ohne historische Originale umzuschreiben."""
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email import policy
from email.utils import format_datetime
from zoneinfo import ZoneInfo
import importlib.util, json, hashlib, os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
CASE=ROOT/'testakten/bauwirtschaft-hoai-buergerhaus-einbeck'
QA=Path(os.environ.get('EINBECK_VERTIEFUNG_QA','/tmp/bauwirtschaft-vertiefung-20261006/einbeck'))
FILES={45:'45_2026-09-22_Nachlieferung_Projektunterlagen.docx',46:'46_2026-09-23_Erlaeuterung_Aufmass_AM07.docx',47:'47_2026-09-24_Geraetebereitschaft_Kalkulationsnotiz.docx',48:'48_2026-09-25_Aufmass_und_Rechnungsabgleich.xlsx',49:'49_2026-09-28_Rueckfragen_Geraetebereitschaft.eml',50:'50_2026-09-29_Hauswart_Betriebsbeobachtungen.docx',51:'51_2026-09-30_Angebot_Anschlussuntersuchung.docx',52:'52_2026-10-01_Leistungsnachweise_Hartwig.docx',53:'53_2026-10-02_Maengel_und_Befundverfolgung.xlsx',54:'54_2026-10-05_Honorar_und_Belegstand.xlsx',55:'55_2026-10-06_Vorstand_Untersuchungsumfang.eml'}

def doc(n,issuer,recipient,title,parts,signer):
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.1);s.right_margin=Cm(2.1)
    for st in d.styles:
        if st.type==1:
            st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
            fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if key.endswith('Theme'):del fonts.attrib[key]
            for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Times New Roman')
            st.paragraph_format.space_after=Pt(7)
            st.paragraph_format.line_spacing=1.05
            for border in st.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    d.styles['Title'].font.size=Pt(16)
    d.styles['Heading 1'].font.size=Pt(12);d.styles['Heading 1'].font.bold=True
    d.styles['Heading 1'].paragraph_format.space_before=Pt(10);d.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    p=d.core_properties;p.author='Klotzkette';p.last_modified_by='Klotzkette';p.title=title;p.language='de-DE'
    p.created=datetime.fromisoformat(FILES[n][3:13]);p.modified=p.created
    d.add_paragraph(issuer).runs[0].bold=True
    d.add_paragraph('Bürgerhaus Leinewinkel · Projekt BH-21-04\n'+recipient)
    d.add_paragraph('Einbeck, '+datetime.fromisoformat(FILES[n][3:13]).strftime('%d.%m.%Y'))
    d.add_paragraph(title,'Title')
    for item in parts:
        if isinstance(item,str):
            p=d.add_paragraph(item);p.paragraph_format.widow_control=True
        elif item[0]=='h':
            break_before={46:'3 ',47:'2 ',50:'3 ',51:'3 ',52:'3 '}.get(n)
            if break_before and item[1].startswith(break_before):d.add_page_break()
            d.add_paragraph(item[1],'Heading 1')
        elif item[0]=='page':d.add_page_break()
        elif item[0]=='table':
            t=d.add_table(rows=0,cols=len(item[1][0]));t.autofit=False
            for idx,row in enumerate(item[1]):
                cells=t.add_row().cells
                for c,txt in zip(cells,row):
                    c.text=str(txt)
                    tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
                    for edge in ('top','left','bottom','right'):
                        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
                    tcpr.append(b)
                    for p in c.paragraphs:
                        p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
                        for r in p.runs:r.bold=idx==0;r.font.size=Pt(11)
                trpr=t.rows[idx]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if idx==0:trpr.append(OxmlElement('w:tblHeader'))
            if len(item)>2:
                for col,w in zip(t.columns,item[2]):col.width=Cm(w)
                for row in t.rows:
                    for c,w in zip(row.cells,item[2]):c.width=Cm(w)
    p=d.add_paragraph('gez. '+signer);p.paragraph_format.keep_with_next=True
    d.add_paragraph('Bezug: Die genannten Nummern bezeichnen eigenständige Unterlagen der Projektmappe. Anlagen: keine.')
    footer=s.footer.paragraphs[0];footer.add_run('BH-21-04 · '+str(n)+' · Seite ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    d.save(CASE/FILES[n])

def eml(n,sender,recipient,subject,body,refs=None):
    m=EmailMessage(policy=policy.SMTP);date=datetime.fromisoformat(FILES[n][3:13]+'T10:20:00').replace(tzinfo=ZoneInfo('Europe/Berlin'))
    m['From']=sender;m['To']=recipient;m['Date']=format_datetime(date);m['Subject']=subject
    m['Message-ID']=f'<BH-21-04-{n}@buergerhaus-leinewinkel.example>';m['Return-Path']='<ablage@buergerhaus-leinewinkel.example>'
    m['Received']=f'from ablage.buergerhaus-leinewinkel.example by archiv.buergerhaus-leinewinkel.example; {format_datetime(date)}'
    if refs:m['References']=refs
    m.set_content(body+'\n\nAnlagen: keine. Die genannten Dateien liegen als eigenständige Unterlagen in der Projektmappe.\n',charset='utf-8')
    (CASE/FILES[n]).write_bytes(m.as_bytes())

def main():
    QA.mkdir(parents=True,exist_ok=True)
    baseline={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in CASE.iterdir() if p.name[:2].isdigit() and int(p.name[:2])<=44}
    (QA/'baseline-originale.json').write_text(json.dumps(baseline,indent=2,ensure_ascii=False))
    doc(45,'Hartwig Architektur · Lena Hartwig','An Maren Döring und Uwe Feldmann, Vorstand','Nachlieferung der Projektunterlagen',[
        'Sehr geehrte Frau Döring, sehr geehrter Herr Feldmann, Sie baten im Gespräch vom 21.09.2026 darum, die Mengenänderungen und die noch offenen Leistungen anhand der Projektmappe erläutern zu können. Wir liefern dazu ergänzende Erläuterungen nach. Der Stand der ursprünglichen Kostenfeststellung vom 08.04.2024 bleibt bestehen.',
        ('h','1 Gegenstand der Nachlieferung'),
        'Die Dateien 09, 19, 24 und 37 sind abgeschlossene Arbeitsstände aus Planung und Bauausführung. Sie werden nicht rückwirkend geändert. Insbesondere ist der in Datei 37 ausgewiesene Betrag von 52.543,02 EUR brutto nach Abschlägen ein damaliger Zahlungsvorschlag. Die Unterlagen enthalten keinen Kontoauszug über eine spätere Schlusszahlung. Eine heute noch offene Geldforderung lässt sich daraus allein nicht bestimmen.',
        'Für die Aufmaßbesprechung haben wir AM-07, das Angebot Leinebau und die Schlussrechnung nebeneinandergelegt. Bei der Rampe lässt sich der zusätzliche Anschlussstreifen aus dem gemeinsam festgehaltenen Rechenweg wiederfinden. Beim Rinnengefälle fehlt dagegen weiterhin ein unterschriebenes Nivellementblatt. Eine feststehende Abrechnungslänge ist kein Nachweis der Entwässerungsfunktion.',
        ('h','2 Fragen aus dem Vorstandsgespräch'),
        'Herr Feldmann fragte, weshalb der Fundamentaushub um sechs Kubikmeter gestiegen sei, obwohl bereits vier Kubikmeter Leitungsgraben beauftragt worden seien. Nach AM-07 sind dies verschiedene Mengen. Wir werden die Trennung in der ergänzenden Mengenrechnung sichtbar halten. Für eine Prüfung der verdeckten Geometrie wäre die ursprüngliche Aufnahme vor Betonage erforderlich; diese liegt in der hier zusammengestellten digitalen Mappe nicht als unterschriebenes Einzelblatt vor.',
        'Frau Döring fragte außerdem, ob mit der Begehung vom 17.09.2026 die Objektbetreuung erledigt sei. Unser Vertrag sieht weitere Betreuung und spätere Kontrollen vor. Der aktuelle Befund ist noch nicht geklärt. Wir werden deshalb die seit 2024 dokumentierten Leistungen und die offenen Arbeitsschritte getrennt ausweisen. Eine Abnahmeerklärung für unsere Leistungen haben wir in dieser Mappe weiterhin nicht gefunden.',
        ('h','3 Weiteres Vorgehen'),
        'Leinebau wurde um eine Erläuterung des Tagessatzes und des tatsächlichen Geräteeinsatzes am 06. und 07.09.2023 gebeten. Der Hauswart stellt seine aktuellen Betriebsbeobachtungen zusammen. Für eine Untersuchung von Rinne und Sockel soll zunächst ein begrenzter Vorschlag mit Preis, Teilnehmern und ausdrücklich bezeichneten Eingriffen vorliegen. Der Vorstand entscheidet anschließend über den Umfang.',
        'Die Nachlieferung entsteht im September und Oktober 2026. Spätere Erinnerungen werden als solche bezeichnet. Sie erhalten weder das Datum noch den Beweiswert einer während der Ausführung erstellten Messurkunde.'], 'Lena Hartwig, Architektin')
    doc(46,'Hartwig Architektur · Lena Hartwig','An Torsten Albrecht und den Vorstand','Erläuterung des Aufmaßes AM 07',[
        'Sehr geehrter Herr Albrecht, für die Rückfragen des Vorstands erläutern wir die am 16.02.2024 in AM-07 festgehaltenen Schlussmengen. Diese Erläuterung ist eine Übertragung aus Datei 35. Sie dokumentiert keine neue gemeinsame Vermessung und ergänzt keine fehlende Unterschrift unter einem ursprünglichen Messblatt.',
        ('h','1 Mengen mit einfachen Flächenbezügen'),
        'Der selektive Rückbau betrifft den Bestand mit 18,00 m mal 12,00 m und damit 216,00 m². Die Bodenplatte und das Dach des Anbaus sind jeweils mit 10,00 m mal 8,00 m und damit 80,00 m² erfasst. Bei der Rampe ergeben 12,00 m Lauflänge mal 1,50 m Breite 18,00 m². Hinzu kommen das Zwischenpodest mit 2,25 m² und der in AM-07 gesondert beschriebene Anschlussstreifen mit 1,20 m². Zusammen sind dies 21,45 m² statt des ursprünglichen Ansatzes von 20,25 m².',
        'Für den Anschlussstreifen enthält AM-07 die Fläche, aber keine getrennten Kantenlängen. Die neue Rechentabelle übernimmt daher 1,20 m² als dokumentierte Zusatzfläche und erfindet keinen nachträglich bemaßten Plan. Bei der Rinne werden 9,00 m Grundlänge und 0,40 m angepasstes Endstück zu 9,40 m addiert.',
        ('h','2 Mengen verdeckter oder ergänzter Leistungen'),
        ('table',[['Position','Übertrag aus AM-07','Verbleibende Nachweisfrage'],['2.3 Aushub','60,00 + 6,00 = 66,00 m³','Aufnahme vor Betonage vom 20.09.2023 nicht als Einzelblatt in dieser Mappe.'],['2.4 Fundamente','24,00 + 1,20 = 25,20 m³','Verbreiterung übernommen, keine neue geometrische Vermessung.'],['2.7 Sockelabdichtung','36,00 + 6,00 = 42,00 m²','Feststellung vor Verfüllung am 03.10.2023 in AM-07 referenziert.'],['2.8 Innenputz','144,00 + 6,00 = 150,00 m²','Zusatzfläche übernommen; keine flächige Feuchtemessung daraus ableitbar.']],[3.1,4.3,9.4]),
        ('h','3 Trennung vom Nachtrag'),
        'Die 4,00 m³ Graben aus N01.2 betreffen die Leitungsumlegung und werden nur im Nachtragsblock gerechnet. Die 66,00 m³ der Position 2.3 enthalten diese Menge nach AM-07 ausdrücklich nicht. Würden beide Mengen in Position 2.3 zusammengeführt und N01.2 unverändert belassen, entstünde eine doppelte Mengenbasis. Die Nachtragsleitung beträgt 14,00 m, der Abtransport eine Pauschale und die Anschlussarbeit 16,00 Stunden.',
        'Für N01.5 gibt es keine gemeinsame Mengenbestätigung. Zwei Tage zu 620,00 EUR sind die Behauptung von Leinebau. Der Auftrag vom 15.09.2023 ließ gerade diese Position offen. Die ergänzende Rechnung kann den geltend gemachten und einen frei einzugebenden Prüfanteil gegenüberstellen; sie kann aus der Zahl der Tage keinen Anspruch ableiten.',
        'Bitte benennen Sie die Ablagestelle der ursprünglichen verdeckten Aufnahmen. Bis zu deren Eingang bleibt die Erläuterung auf AM-07 beschränkt. Die Mengenfeststellung sagt nichts Abschließendes über Mängelfreiheit, Höhenanschluss oder eine Zahlung aus.'], 'Lena Hartwig, Architektin')
    doc(47,'Leinebau Hochbau GmbH · Torsten Albrecht','An Hartwig Architektur und den Vorstand','Erläuterung der Gerätebereitschaft N01 5',[
        'Sehr geehrte Frau Hartwig, wir halten den in unserer Schlussrechnung LB-240301 enthaltenen Ansatz von zwei Tagen Gerätebereitschaft zu jeweils 620,00 EUR netto aufrecht. Auf Ihre Nachfrage erläutern wir die Kalkulation und unsere Erinnerung an den Ablauf. Dies ist keine neue Rechnung und kein beiderseits unterzeichneter Tagesnachweis.',
        ('h','1 Zusammensetzung unseres Ansatzes'),
        ('table',[['Bestandteil je Tag','Ansatz netto','Grundlage unserer Erklärung'],['Gerätevorhaltung Kompaktbagger','320,00 EUR','Kalkulatorischer Tagessatz unseres eigenen Geräts; kein Mietbeleg eines Dritten.'],['Fahrerbereitschaft','8 h × 32,50 EUR = 260,00 EUR','Interner Verrechnungssatz einschließlich Personalnebenkosten.'],['Baustellenorganisation','40,00 EUR','Interner Ansatz für Disposition und Sicherung.'],['Summe je Tag','620,00 EUR','Keine gesonderte Umsatzsteuer in den Komponenten.']],[4.2,3.7,8.9]),
        'Der Betrag von 320,00 EUR ist nicht der Anschaffungspreis und auch keine Stundenabrechnung nach Motorlaufzeit. Eine Aufteilung in Abschreibung, Verzinsung und tatsächliche Reparaturkosten liegt der Projektmappe nicht bei. Wir können dazu eine gesonderte Aufstellung anfordern. Die 40,00 EUR sind nicht durch eine zusätzliche Fremdrechnung belegt.',
        ('h','2 Erinnerung an den 06 und 07 September 2023'),
        'Am 06.09.2023 wurde im Bestandsbereich weiter zurückgebaut. Dies bestreiten wir nicht. Der für die Fundamentachse vorgesehene Bagger blieb nach unserer Erinnerung am offenen Graben. Der Fahrer half zeitweise beim Absperren und bei der Sicherung des Leitungsfundes. Einen damals gegengezeichneten Stundenzettel, der diese Arbeiten minutengenau von Wartezeit trennt, können wir heute nicht vorlegen.',
        'Am 07.09.2023 warteten wir auf die Klärung, ob der Pumpenstrang in Betrieb bleiben musste. Es wurde Material umgelagert. Herr Albrecht erinnert eine Unterbrechung des Aushubs bis zur Entscheidung über den neuen Verlauf; Beginn und Ende einzelner Bewegungen des Baggers wurden im allgemeinen Bautagebuch nicht erfasst. Die Zuordnung des Strangs erfolgte erst am 11.09.2023. Daraus folgt aus unserer Sicht aber nicht, dass für jeden Zwischentag Gerätebereitschaft abgerechnet wurde.',
        ('h','3 Abgrenzung und offene Unterlagen'),
        'Wir haben nur den 06. und 07.09.2023 angesetzt. Der Fahrer war nicht durchgehend untätig. Wir sehen Sicherungsarbeiten als Teil der angeordneten Bereithaltung; eine gesonderte Vergütungsvereinbarung hierzu können wir nicht beifügen. Über den Einwand, dass Vorhaltung oder Personal bereits in der Baustelleneinrichtung beziehungsweise in ausgeführten Positionen enthalten sein könnten, besteht keine Einigung.',
        'Eine von uns vorgeschlagene hälftige Anerkennung soll nicht stillschweigend in die Schlussprüfung eingetragen werden. Bitte teilen Sie uns mit, welche Unterlagen Sie für eine weitere Prüfung benötigen. Unsere Erläuterung ersetzt weder den Auftrag N01.1 bis N01.4 noch dessen ausdrücklichen Vorbehalt zu N01.5.'], 'Torsten Albrecht, Bauleiter')
    eml(49,'Lena Hartwig <projekt@hartwig-architektur.example>','Torsten Albrecht <bauleitung@leinebau-hochbau.example>','BH-21-04 Gerätebereitschaft und ergänzender Rechnungsabgleich',
        'Sehr geehrter Herr Albrecht,\n\nvielen Dank für Ihre Erläuterung vom 24.09.2026 (Datei 47). Die drei Bestandteile ergeben rechnerisch 620,00 EUR je Tag. Die Rechnung beantwortet noch nicht, welcher Anteil zusätzlich entstanden ist und welcher Teil auf tatsächlich geleistete Sicherungs- oder Rückbauarbeiten entfällt. Bitte reichen Sie die vorhandenen Lohnstundenaufzeichnungen, den Geräteeinsatzkalender und die Grundlage des Vorhaltungssatzes nach. Nicht vorhandene Aufzeichnungen kennzeichnen Sie bitte ausdrücklich.\n\nIn der Datei 48 steht die Ausgangsprüfung weiterhin bei 0,00 EUR für N01.5. Das bedeutet, dass der Betrag bisher nicht freigegeben ist. Es ist keine neue abschließende Zurückweisung. Eine Eingabe von 50 Prozent ergäbe 620,00 EUR netto und 737,80 EUR brutto zusätzlich; dies wäre lediglich ein Rechenszenario. Der Vorstand hat einen solchen Vergleich nicht beschlossen.\n\nDer Grundauftrag ergibt nach AM-07 weiterhin 110.447,80 EUR netto. Mit den beauftragten Nachtragsleistungen sind es 114.153,80 EUR netto. Der historische Vorschlag nach Abschlägen beträgt 52.543,02 EUR brutto. Bitte nennen Sie den seit März 2024 tatsächlich eingegangenen Schlusszahlungsbetrag und senden Sie einen zuordenbaren Zahlungsnachweis. Unsere Projektmappe enthält dazu keine Bestätigung.\n\nDie Mängeluntersuchung führen wir als eigenen Vorgang. Die Höhe des Zahlungsansatzes legt die Ursache der aktuellen Erscheinungen nicht fest.\n\nMit freundlichen Grüßen\nLena Hartwig')
    doc(50,'Bürgerhaus Leinewinkel e.V. · Ralf Winter, Hauswart','An Maren Döring und Lena Hartwig','Beobachtungen an Rinne und Saalwand',[
        'Frau Döring bat mich, meine Beobachtungen nach der Begehung vom 17.09.2026 schriftlich festzuhalten. Ich habe den Putz nicht geöffnet, keine elektrische Feuchtemessung vorgenommen und die Grundstücksleitung nicht gespült. Die nachstehenden Zeiten stammen aus meinen aktuellen Kalendereinträgen. Meine Erinnerung an die Reinigung vom 03.09.2026 habe ich davon getrennt.',
        ('h','1 Rückblick auf die erste Meldung'),
        'Am 03.09.2026 hatte ich Blätter und kleine Zweige vom Rost entfernt. In meinem Kalender steht nur „Eingang gekehrt, Rinne sauber“. Ob der Ablaufstutzen damals frei war, habe ich nicht geprüft. Am 04.09.2026 sah ich nach einem Regenschauer Wasser vor T-03. Niederschlagsmenge und Dauer habe ich nicht gemessen. Das Foto der Wand vom 07.09.2026 ist Datei 42; es zeigt keine Aufnahme des Wassers vor der Tür.',
        ('h','2 Beobachtungen seit der gemeinsamen Begehung'),
        ('table',[['Zeitpunkt','Beobachtung','Was ich tatsächlich getan habe'],['21.09.2026, 08:10','Rost fest; wenige Blätter im Rinnenkasten. Keine Pfütze.','Blätter von Hand entfernt, Anschluss nicht geöffnet.'],['24.09.2026, 17:35','Nach einem Schauer Wasser am seitlichen Anschluss. Mit Zollstock etwa 4 mm tief.','Zugang vorübergehend abgesperrt. Kein Foto gefertigt.'],['24.09.2026, 18:10','Wasserfläche kleiner; schmaler Rand noch sichtbar.','Keine Reinigung zwischen beiden Beobachtungen.'],['28.09.2026, 09:20','Wand links vom Nordfenster unverändert fleckig; loser Farbspan am Boden.','Span beim Kehren entsorgt; keine Probe zurückgelegt.'],['29.09.2026, 08:45','Kein Wasser vor T-03; Rost lässt sich nicht klappernd bewegen.','Sichtkontrolle vor Öffnung des Hauses.']],[3.0,6.3,7.5]),
        ('h','3 Grenzen meiner Angaben'),
        'Die Tiefe vom 24.09.2026 habe ich an einer erreichbaren Stelle abgelesen, nicht als vermessene Höhenlage des Bauteils. Sie lässt sich nicht unmittelbar mit dem Höhenversatz von etwa 6 mm aus der Begehung vergleichen. Der kleine Wasserrand war nach meinem zweiten Rundgang noch sichtbar; wann er vollständig verschwand, weiß ich nicht.',
        'Die Verfärbung erschien mir an denselben Stellen wie am 07.09.2026. Ich habe die Ränder nicht markiert und keinen Maßstab angelegt. Deshalb kann ich weder eine sichere Ausbreitung noch eine Abnahme angeben. Die Türen bleiben im üblichen Betrieb nutzbar; bei Wasser vor dem Eingang leite ich Besucher zum freigehaltenen anderen Zugang und informiere den Vorstand.',
        'Für einen Untersuchungstermin bin ich am 09.10.2026 vormittags verfügbar. Dieser Termin ist von der Firma noch nicht bestätigt. Ich werde vorher keine Bauteile entfernen, damit die Beteiligten den Ausgangszustand sehen können.'], 'Ralf Winter, Hauswart')
    doc(51,'Leinebau Hochbau GmbH · Torsten Albrecht','An Bürgerhaus Leinewinkel e.V., Vorstand','Angebot zur Untersuchung des Rinnenanschlusses',[
        'Sehr geehrte Frau Döring, aufgrund der Begehung vom 17.09.2026 und Ihrer Bitte um einen begrenzten Untersuchungsvorschlag bieten wir die nachstehenden Arbeiten an. Das Angebot betrifft die Ermittlung des Zustands. Ein Sanierungserfolg und die Kostentragung für eine spätere Mangelbeseitigung werden damit nicht vereinbart.',
        ('h','1 Untersuchung ohne Bauteilöffnung'),
        'Zum Pauschalpreis von 620,00 EUR netto nehmen wir die zugänglichen Höhen am letzten Rinnenelement, der Türschwelle und dem angrenzenden Belag mit Bezug auf einen gemeinsam festgelegten Punkt auf. Wir öffnen den herausnehmbaren Rost, dokumentieren vorhandenes Sediment und prüfen den zugänglichen Ablauf mit einer kleinen Kamera, soweit der Verlauf die Einführung erlaubt. Erst nach der Aufnahme wird ein begrenzter Wasserversuch mit protokollierter Menge durchgeführt.',
        'Enthalten sind zwei Mitarbeiter für zusammen sechs Arbeitsstunden, Anfahrt und ein Protokoll mit Messpunktskizze. Eine flächige Reinigung, Druckspülung oder Demontage der Grundstücksleitung ist nicht enthalten. Ist die Leitung für die Kamera nicht passierbar, wird die Stelle dokumentiert und vor einem weiteren Eingriff Rücksprache gehalten. Aus dem bloßen Abbruch einer Kamerafahrt wird keine Ursache abgeleitet.',
        ('h','2 Wahlposition zur örtlichen Öffnung'),
        'Eine zusätzliche Öffnung am Übergang Rinne und Grundstücksleitung bieten wir als Wahlposition zu 980,00 EUR netto an. Sie umfasst die Aufnahme von bis zu 1,00 m² Pflaster, vorsichtige Freilegung bis 0,60 m Tiefe und die anschließende Wiederherstellung desselben Bereichs. Eine Öffnung des Saalwandputzes oder der Sockelabdichtung ist davon nicht umfasst. Für einen solchen Eingriff benötigen wir einen gesonderten Umfang und eine abgestimmte Sicherung der Proben.',
        'Die Wahlposition wird nur nach gesonderter schriftlicher Beauftragung ausgeführt. Bei unerwarteten Leitungen oder einer erforderlichen größeren Öffnung unterbrechen wir die Arbeiten. Die erstmalige Wiederherstellung im angebotenen Umfang ist enthalten; zusätzliche Sanierungsarbeiten sind es nicht.',
        ('h','3 Preis und Abstimmung'),
        ('table',[['Umfang','Netto EUR','Umsatzsteuer EUR','Brutto EUR'],['Nur Untersuchung nach Abschnitt 1','620,00','117,80','737,80'],['Zusätzliche Wahlposition nach Abschnitt 2','980,00','186,20','1.166,20'],['Beide Abschnitte zusammen','1.600,00','304,00','1.904,00']],[7.0,3.3,3.3,3.2]),
        'Wir schlagen den 09.10.2026 um 09:00 Uhr vor. Der Termin ist bis 06.10.2026 freizuhalten und danach erneut abzustimmen. Frau Hartwig und der Vertreter der Außenanlagen sollen Gelegenheit zur Teilnahme erhalten. Wir bitten den Verein um die Koordination. Bis zum heutigen Datum liegt uns keine Beauftragung vor.',
        'Die Mitwirkung an der Untersuchung ist keine Erklärung zur Verursachung oder zur Übernahme aller Kosten. Wir halten unseren Hinweis auf die Bestandsfeuchte aufrecht. Zugleich wollen wir den Rinnenanschluss gemeinsam nachvollziehbar aufnehmen. Dieses Angebot gilt bis 14.10.2026.'], 'Torsten Albrecht, Bauleiter')
    doc(52,'Hartwig Architektur · Lena Hartwig','An Bürgerhaus Leinewinkel e.V., Vorstand','Leistungsnachweise und offene Unterlagen',[
        'Sehr geehrte Frau Döring, Sie baten um eine Erklärung, welche Arbeitsergebnisse hinter den Phasenbezeichnungen unseres Vertrags stehen. Wir ordnen die vorhandenen Dateien den tatsächlich bearbeiteten Themen zu. Das ist unsere Darstellung des Leistungsstands am 01.10.2026, keine neue Honorarrechnung und keine Erklärung Ihres Vorstands zur Abnahme.',
        ('h','1 Planung bis zur Genehmigung'),
        'Aus dem Bedarfsbeschluss 01 ergaben sich Saal, Besprechungsraum und barrierearmer Zugang als Planungsaufgabe. Das Grundlagengespräch 03 und die Bestandsaufnahme 04 hielten die Untersuchungsstellen fest. Die Bestandsarbeiten sind nach Vertrag gesondert pauschal vergütet; sie werden nicht nochmals als Prozentanteil der Grundlagenermittlung berechnet.',
        'In der Vorplanung wurden Süd- und Nordanbau anhand gleicher Nutzungsanforderungen verglichen (05). Die Entscheidung 06 gab den Nordanbau vor. Entwurf 07, Objektbeschreibung 08, Kostenberechnung 09 und Fachplanerabstimmung 10 bilden den anschließenden Stand. Die Kostenberechnung verwendete eine vorsichtige Steuerpauschale; die Kostenfeststellung 37 führt Gebühren später ohne Umsatzsteuer. Dieser Unterschied ist kein Wechsel der vertraglichen Honorarpauschale.',
        'Die Genehmigungsplanung umfasst Antrag 11 und Bauvorlagen 12 bis 14. Auf die Nachforderung 15 wurde mit der Brandschutzergänzung 16 reagiert; 17 enthält den Genehmigungsbescheid. Der spätere Wunsch nach Gymnastiknutzung aus 29 wurde nicht als genehmigte Umplanung ausgeführt.',
        ('h','2 Ausführung und Vergabe'),
        'Nach dem Abruf 18 wurden die Ausführungsstände AP-01 und D-12 (20 und 21) sowie der Rahmenterminplan 19 geführt. Die Portalansicht 22 unterscheidet die Freigaben. Der Koordinationsstand TGA-04 ersetzt nach unserer Ablage nicht D-12. Ob die abweichende Höheninformation bei der Ausführung hinreichend eindeutig aufgelöst wurde, soll anhand des tatsächlichen Anschlusses geprüft werden.',
        'Das Leistungsverzeichnis 23 und die bepreiste Fassung 24 enthalten dieselben zwölf Positionen. Die Angebote 25 und 26 sind im Preisspiegel 27 positionsbezogen verglichen. 28 dokumentiert den Zuschlag. Die zusätzliche Pumpenleitung wurde über Nachtragsangebot 33 und Teilauftrag 34 behandelt. Die ausgeschlossene Gerätebereitschaft wird nicht durch die technische Prüfung der Leitungsarbeiten zu einem Auftrag.',
        ('h','3 Überwachung und Betreuung'),
        'Bautagebuch 30, Aufmaß 35, Schlussrechnung 36 und Prüfung 37 zeigen den Bau- und Abrechnungsstand. Die gemeinsame Mengentabelle ersetzt keine Messung des Rinnengefälles. Zur Nachkontrolle von M-01 und M-02 enthält 39 einen Eintrag; ein eigenes Nachkontrollblatt vom 28.03.2024 ist in dieser Mappe nicht vorhanden. Die Abnahme 38 betrifft Los 1. Die gesonderte Abnahme unserer Planungs- und Überwachungsleistungen war bei Übergabe 39 noch vorgesehen.',
        'Die Dateien 41 bis 44 dokumentieren Meldung, Foto, Begehung und laufende Betreuung. Die fachliche Zuordnung der heutigen Erscheinungen ist noch offen. Eine weitere Begehung vor den gewerkeweisen Kontrollterminen bleibt zu organisieren. Die Pauschale von 960,00 EUR netto für Phase 9 wird im bisherigen Kostenstand deshalb gesondert geführt; aus der einzelnen Begehung folgt kein automatischer voller Erledigungsgrad.',
        ('h','4 Bitte um Abgleich'),
        'Bitte teilen Sie uns mit, ob die gesonderte Abnahmeerklärung oder weitere Einzelblätter in Ihrem Papierarchiv liegen. Unsere ergänzende Arbeitsmappe wird Honorarzuordnung, rechnerisch angesetzten Leistungsanteil und Belegstatus getrennt zeigen. Prozente sind Rechengrößen aus dem individuell vereinbarten Vertrag. Sie ersetzen weder eine Prüfung einzelner Leistungen noch eine Entscheidung über Fälligkeit, Mängel oder Zahlung.'], 'Lena Hartwig, Architektin')
    eml(55,'Maren Döring <vorstand@buergerhaus-leinewinkel.example>','Lena Hartwig <projekt@hartwig-architektur.example>','BH-21-04 Umfang der Untersuchung und noch fehlende Unterlagen',
        'Sehr geehrte Frau Hartwig,\n\nHerr Feldmann und ich haben Ihre Ergänzungen und das Angebot von Leinebau gelesen. Wir haben heute weder die Wahlposition zur Öffnung noch die Pauschale für die Untersuchung beauftragt. Bitte klären Sie zunächst, wer die Grundstücksleitung aus dem Los Außenanlagen vertritt und ob dessen Teilnahme zusätzliche Kosten verursacht. Den 09.10.2026 können wir deshalb noch nicht als gemeinsamen Termin bestätigen.\n\nUns ist nach Ihrer Tabelle klarer, weshalb die vier Kubikmeter Leitungsgraben nicht nochmals in den 66 Kubikmetern Fundamentaushub stehen. Das ursprüngliche Blatt vor Betonage haben wir im Vereinsordner noch nicht gefunden. Herr Feldmann sucht am Donnerstag weiter. Die Mengen unterschreiben wir heute nicht neu.\n\nZur Gerätebereitschaft enthält die Erklärung des Bauleiters weiterhin keine getrennten Stunden. Wir stimmen dem angesprochenen Ansatz von 50 Prozent derzeit nicht zu. Die Suche nach den Kontoauszügen über die Schlusszahlung läuft bei unserer früheren Kassenwartin. Bitte behandeln Sie den Betrag aus 2024 daher weiterhin als damaligen Prüfungsvorschlag, nicht als festgestellten heutigen Rückstand.\n\nIn der Honorarübersicht ist bei Phase 8 eine offene Belegfrage vermerkt, obwohl der rechnerische Anteil vollständig angesetzt ist. Wir möchten im Gespräch verstehen, welche Unterlagen dazu fehlen und ob der Eintrag im Übergabeprotokoll zur Nachkontrolle ausreicht. Die Existenz der Tabelle bedeutet keine Zustimmung zu einem Leistungsgrad. Die gesonderte Abnahme der Architektenleistungen ist im heute durchgesehenen Ordner nicht auffindbar.\n\nHerr Winter hält den Wandbereich frei und protokolliert neue Wasserbeobachtungen. Bitte schicken Sie uns nach der Abstimmung einen klar begrenzten Untersuchungsvorschlag mit Gesamtpreis und Teilnehmern. Eine Bauteilöffnung soll erst nach einer gesonderten Entscheidung erfolgen.\n\nMit freundlichen Grüßen\nMaren Döring\nfür den Vorstand des Bürgerhauses Leinewinkel e.V.')
    (QA/'neue-dateien.json').write_text(json.dumps(FILES,indent=2,ensure_ascii=False))
    print(json.dumps({'docx':6,'eml':2,'new_files':FILES},ensure_ascii=False))

if __name__=='__main__':main()
