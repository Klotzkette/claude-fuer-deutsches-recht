#!/usr/bin/env python3
"""Prüft CISG-Paket, Quellenpflege und die drei konsistenten Aktenfassungen."""

import argparse
import csv
from datetime import datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import tempfile
import zipfile

from docx import Document
from pypdf import PdfReader
import yaml

from testakte_disclaimer import NOTICE_DE, NOTICE_EN, NOTICE_BYTES

ROOT = Path(__file__).resolve().parents[1]
NAME = "un-kaufrecht-cisg"
PLUGIN = ROOT / NAME


def require(value, message):
    if not value:
        raise ValueError(message)


def module(name):
    spec = importlib.util.spec_from_file_location(name.replace("-", "_"), ROOT / "scripts" / f"{name}.py")
    loaded = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(loaded)
    return loaded


def pages(data):
    reader = PdfReader(io.BytesIO(data))
    require(reader.pages, "Leeres PDF")
    result = [re.sub(r"\s+", " ", page.extract_text()).strip() for page in reader.pages]
    for text in result:
        require(len(text) > 30, "Leere oder verstümmelte PDF-Seite")
        require(NOTICE_DE not in text and NOTICE_EN not in text, "Warntext im PDF statt auf Downloadseite")
    return result


def check(dist=None):
    downloads = module("validate-testakten-readme-downloads")
    skills = sorted((PLUGIN / "skills").glob("*/SKILL.md"))
    require(len(skills) == 11, "Zehn Fachskills und ein Hauptskill")
    names = {path.parent.name for path in skills}
    for path in skills:
        text = path.read_text(encoding="utf-8")
        require(text.startswith("---\n"), f"Frontmatter fehlt: {path}")
        front = yaml.safe_load(text.split("---", 2)[1])
        require(set(front) == {"name", "description"}, f"Frontmatter-Felder: {path}")
        require(front["name"] == path.parent.name, f"Skillname: {path}")
        require(80 <= len(front["description"]) <= 1024, f"Beschreibungslänge: {path}")
        require(len(text.splitlines()) <= 500 and len(text) >= 1800, f"Skillumfang: {path}")
        require("§" not in text and "**" not in text, f"Skillstil: {path}")
        for n in range(1, 7):
            require(f"## 1.{n}." in text, f"Fachabschnitt {n} fehlt: {path}")
        for target in re.findall(r"`(\.\./[^`]+)`", text):
            require((path.parent / target).is_file(), f"Tote Referenz: {path}/{target}")
        for target in re.findall(r"\[[^\]]+\]\(([^)]+)\)", text):
            if target.startswith(("https://", "#")):
                continue
            resolved = (path.parent / target.split("#", 1)[0]).resolve()
            require(resolved.is_relative_to(PLUGIN.resolve()) and resolved.is_file(),
                    f"Referenz fehlt im installierbaren Paket: {path}/{target}")
        require("CISG" in text and "Artikel" in text, f"Fachlicher Normbezug: {path}")
    legal_guards = {
        "cisg-nacherfuellung-aufhebung-waehlen": ["Nichtlieferung", "Artikel 26", "keine pauschale Pflicht zu zwei Versuchen", "keine Sachentscheidung"],
        "cisg-gefahr-transport-abgrenzen": ["nicht automatisch", "Absatz 2", "Gegenbeweis", "21-Tage", "Artikel 32", "CISG-Rüge an Verkäufer zusätzlich"],
        "cisg-rueckabwicklung-fristen-sichern": ["keine Verjährungsnorm", "Artikel 82", "Artikel 3 UNWaVtrÜbkG", "keine Hemmung"],
        "cisg-schadensersatz-zinsen-berechnen": ["Ausfallumsatz ist nicht entgangener Gewinn", "nennt aber keinen Zinssatz", "nicht rückwirkend", "keine automatische Vertragsaufhebung"],
    }
    for slug, phrases in legal_guards.items():
        body = (PLUGIN / "skills" / slug / "SKILL.md").read_text()
        for phrase in phrases:
            require(phrase.casefold() in body.casefold(), f"Fachliche Fehlergrenze fehlt: {slug}: {phrase}")
    bodies = {kind: (PLUGIN / f"{NAME}-{kind}.md").read_text(encoding="utf-8")
              for kind in ("werkstatt", "schnellstart", "hauptproblem")}
    for kind in ("schnellstart", "hauptproblem"):
        body = bodies[kind]
        require(max(len(body), len(body.encode()), len(body.replace("\n", "\r\n").encode())) <= 7500,
                f"{kind}: Grenze einschließlich UTF-8 und CRLF")
        require(max(map(len, body.split("\n\n"))) <= 1200, f"Unlesbarer Absatz: {kind}")
    require(len(bodies["werkstatt"].encode()) >= 25000, "Ausführliche Werkstatt fehlt")
    for kind, text in bodies.items():
        require("§" not in text and "**" not in text, f"Promptstil: {kind}")
        require(all(re.match(r"^#{1,6} \d", line) for line in text.splitlines() if line.startswith("#")),
                f"Dezimale Überschrift: {kind}")
        require("CISG" in text[:1200] and "Artikel" in text[:2000], f"Kein fachlicher Einstieg: {kind}")
    profile = json.loads((ROOT / "quality/evals" / f"{NAME}.json").read_text())
    module("quality_lab").validate_profile(profile, NAME, PLUGIN)
    require({case["target_skill"] for case in profile["cases"]} == names, "Nicht jeder Skill mit eigenem Prüffall abgedeckt")
    require(len(profile["cases"]) >= 11, "Elf fachliche Prüffälle fehlen")
    references = "\n".join(path.read_text() for path in (PLUGIN / "references").glob("*.md"))
    for decision in profile["prompt_editorial_review"]["decisions"]:
        require(decision["url"] in references, f"Unauffindbare Entscheidungsquelle: {decision['citation']}")
        require(decision["limit"] and decision["application"], "Entscheidung ohne Reichweitengrenze")
    errors = []
    count = downloads.validate_download_notices(f"{NAME}/README.md", (PLUGIN / "README.md").read_text(), errors)
    require(count == 3 and not errors, f"Plugin-Downloadhinweise: {errors}")
    cases = json.loads((ROOT / "scripts/data/cisg-akten.json").read_text())
    require(len(cases) == 3 and len({case["slug"] for case in cases}) == 3, "Drei verschiedene Fälle erforderlich")
    module("build-cisg-akten").validate_data(cases)
    total_pages = total_originals = 0
    for case in cases:
        slug = case["slug"]
        require(re.fullmatch(r"[a-z0-9-]{1,64}", slug), "Aktenslug")
        directory = ROOT / "testakten" / slug
        require(12 <= len(case["files"]) <= 20, f"Aktenumfang: {slug}")
        require(downloads.validate_local_readme(slug, directory, errors) == 6 and not errors,
                f"Akten-Downloadstandard: {slug}: {errors}")
        originals, attachments_count = {}, 0
        by_name = {item["name"]: item for item in case["files"]}
        for item in case["files"]:
            path = directory / item["name"]
            require(path.name.isascii() and path.is_file() and path.parent == directory, f"Originalpfad: {path}")
            data = path.read_bytes()
            originals[path.name] = data
            if path.suffix == ".docx":
                doc = Document(io.BytesIO(data))
                body = "\n".join(p.text for p in doc.paragraphs)
                body += "\n".join(c.text for table in doc.tables for row in table.rows for c in row.cells)
                require(len(body) >= 900 and item["title"] in body, f"Unvollständiges Word-Dokument: {path}")
                require(doc.styles["Normal"].font.name == "Times New Roman", f"Dokumentschrift: {path}")
                normalized = re.sub(r"\s+", " ", body)
                require(all(re.sub(r"\s+", " ", paragraph) in normalized for paragraph in item["body"].split("\n\n")),
                        f"Word-Dokument nicht synchron mit Aktenquelle: {path}")
                require(doc.core_properties.author == "Klotzkette", f"Dokumentautor: {path}")
            elif path.suffix == ".eml":
                mail = BytesParser(policy=policy.default).parsebytes(data)
                require(all(mail[h] for h in ("From", "To", "Date", "Subject", "Message-ID")), f"Mailheader: {path}")
                require(not mail.defects, f"Defektes Mailformat: {path}")
                if item.get("reply_to"):
                    reply = BytesParser(policy=policy.default).parsebytes(
                        (directory / item["reply_to"]).read_bytes())
                    for header in ("In-Reply-To", "References"):
                        require(str(mail[header]) == str(reply["Message-ID"]),
                                f"Falscher Nachrichtenbezug: {path}/{header}")
                        raw = next(value for key, value in mail.raw_items() if key == header)
                        require("=?" not in raw and raw.strip().startswith("<"),
                                f"Unzulässig kodierte Nachrichtenkennung: {path}/{header}")
                body = mail.get_body(preferencelist=("plain",)).get_content()
                require(len(body) >= 450, f"Mailstummel: {path}")
                require(body.replace("\r\n", "\n").strip() == item["body"].strip(), f"Mailtext nicht synchron mit Aktenquelle: {path}")
                attached = {p.get_filename(): p.get_payload(decode=True) for p in mail.iter_attachments()}
                require(set(attached) == set(item.get("attachments", [])), f"Mailanlagen fehlen: {path}")
                for name, payload in attached.items():
                    require(payload == (directory / name).read_bytes(), f"Mailanlage veraltet: {name}")
                    dated = datetime.strptime(by_name[name]["date"], "%d.%m.%Y").date()
                    require(dated <= datetime.fromisoformat(item["date"]).date(), f"Mail enthält erst später entstandene Anlage: {path}/{name}")
                attachments_count += len(attached)
            elif path.suffix == ".csv":
                body = data.decode("utf-8-sig")
                rows = list(csv.reader(io.StringIO(body), delimiter=";"))
                require(len(rows) >= 7 and len({len(row) for row in rows}) == 1, f"CSV-Spalten oder Datenumfang: {path}")
                require(rows == item["rows"], f"CSV nicht synchron mit Aktenquelle: {path}")
                if rows[0][-3:] == ["Soll EUR", "Haben EUR", "Offen EUR"]:
                    balance = Decimal(0)
                    for row in rows[1:]:
                        debit, credit, current = (Decimal(value.replace(",", ".")) for value in row[-3:])
                        balance += debit - credit
                        require(balance == current, f"Kontoblatt rechnerisch inkonsistent: {path}/{row[0]}")
            else:
                require(path.suffix == ".txt", f"Unerwartetes Originalformat: {path}")
                body = data.decode("utf-8")
                require(len(body) >= 500, f"Textstummel: {path}")
                require(item["body"] in body, f"Text nicht synchron mit Aktenquelle: {path}")
            require("§" not in body, f"Paragrafenzeichen: {path}")
            require(not any(word in body.casefold() for word in ("musterlösung", "antwortmatrix", "lorem ipsum", "platzhalter", "testakte")),
                    f"Verräterischer Metatext: {path}")
        require(attachments_count >= 1, f"Keine echte Mailanlage: {slug}")
        require({Path(name).suffix for name in originals} >= {".docx", ".eml", ".csv", ".txt"}, f"Fehlender Formatmix: {slug}")
        full = directory / "gesamt-pdf" / f"{slug}_gesamt.pdf"
        full_pages = pages(full.read_bytes())
        total_pages += len(full_pages)
        total_originals += len(originals)
        with zipfile.ZipFile(PLUGIN / "testakte/downloads" / f"testakte-{slug}-einzelpdfs.zip") as archive:
            expected = {Path(name).stem + ".pdf" for name in originals} | {"README.txt"}
            require(set(archive.namelist()) == expected, f"PDF-ZIP enthält Ordner oder falsche Dokumente: {slug}")
            require(archive.read("README.txt") == NOTICE_BYTES and archive.testzip() is None, "PDF-ZIP defekt oder Hinweis fehlt")
            separate = []
            for name in originals:
                separate += pages(archive.read(Path(name).stem + ".pdf"))
            require(separate == full_pages, f"Gesamt-PDF weicht von Einzel-PDFs ab: {slug}")
        with zipfile.ZipFile(PLUGIN / "testakte/downloads" / f"testakte-{slug}.zip") as archive:
            require(set(archive.namelist()) == set(originals) | {"README.txt"}, f"Originalpaket enthält Ordner, Markdown oder doppelte Lesefassung: {slug}")
            require(archive.read("README.txt") == NOTICE_BYTES and archive.testzip() is None, "Originalpaket defekt oder Hinweis fehlt")
            for name, data in originals.items():
                require(archive.read(name) == data, f"Original im Paket veraltet: {name}")
    manifests = [json.loads((PLUGIN / path).read_text()) for path in
                 ("plugin.json", ".claude-plugin/plugin.json", ".codex-plugin/plugin.json")]
    require(len({(value["name"], value["version"], value["description"]) for value in manifests}) == 1, "Manifestkohärenz")
    marketplace = json.loads((ROOT / ".claude-plugin/marketplace.json").read_text())["plugins"]
    matches = [entry for entry in marketplace if entry["name"] == NAME]
    require(len(matches) == 1 and matches[0]["version"] == manifests[0]["version"], "Marketplaceversion")
    if dist:
        require(len(list(dist.iterdir())) == 26, "26 Release-Dateien erforderlich")
        for portable in (False, True):
            prefix = NAME + "/" if portable else ""
            with zipfile.ZipFile(dist / f"{NAME}{'-portable' if portable else ''}.zip") as archive:
                require(archive.testzip() is None, "Plugin-ZIP defekt")
                require(sum(name.endswith("/SKILL.md") for name in archive.namelist()) == 11, "Paket-Skillzahl")
                for name in archive.namelist():
                    relative = name.removeprefix(prefix)
                    source = ROOT / "LICENSE" if relative == "LICENSE" else PLUGIN / relative
                    require(source.is_file() and archive.read(name) == source.read_bytes(), f"Abweichender Paketinhalt: {name}")
                    require(not relative.endswith(("-werkstatt.md", "-schnellstart.md", "-hauptproblem.md", ".pdf", ".zip", ".eml", ".docx")),
                            "Zusatzmaterial im installierbaren Plugin")
        for line in (dist / "checksums-sha256.txt").read_text().splitlines():
            digest, name = line.split("  ", 1)
            require(hashlib.sha256((dist / name).read_bytes()).hexdigest() == digest, f"Prüfsumme: {name}")
        builder = module("build-cisg-release")
        with tempfile.TemporaryDirectory(prefix="cisg-release-test-") as temporary:
            rebuilt = Path(temporary)
            builder.build(rebuilt)
            require(all(path.read_bytes() == (rebuilt / path.name).read_bytes() for path in dist.iterdir()), "Release nicht reproduzierbar")
            try:
                builder.build(rebuilt)
            except ValueError:
                pass
            else:
                raise ValueError("Builder überschreibt vorhandenen Releaseordner")
    print(f"OK: 11 Skills, 3 eigenständige Prompts, 3 Akten, {total_originals} Originale, {total_pages} PDF-Seiten. Technische Regression und redaktionelles Prüfprofil, kein Live-Modelltest.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dist", type=Path)
    check(parser.parse_args().dist)
