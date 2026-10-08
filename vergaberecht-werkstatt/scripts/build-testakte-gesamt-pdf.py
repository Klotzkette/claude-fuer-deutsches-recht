#!/usr/bin/env python3
"""Baut für jede Testakte ein 'gesamt-pdf/<name>_gesamt.pdf', das die
exportfaehigen Aktenstücke (MD/TXT/EML/CSV/XLSX/DOCX/Bilder/PDF) in ein
einziges, sauber gerendertes Dokument mit Dateigrenzen und Seitenzahlen
zusammenfasst.

Aufruf:
  python3 scripts/build-testakte-gesamt-pdf.py                 # alle Testakten
  python3 scripts/build-testakte-gesamt-pdf.py <name1> <name2>  # gezielt
"""

from __future__ import annotations

import argparse
from contextlib import contextmanager
import io
import re
import sys
import csv
import shutil
import tempfile
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path

# Drittabhaengigkeiten
from pypdf import PdfReader, PdfWriter
from reportlab import rl_config
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor, black
from reportlab.platypus import (
    SimpleDocTemplate,
    Image as RLImage,
    Paragraph,
    Spacer,
    PageBreak,
    Table,
    TableStyle,
    HRFlowable,
    Flowable,
    KeepTogether,
)
from reportlab.lib.utils import ImageReader
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from testakte_file_filter import include_in_working_dump

# ReportLab erzeugt sonst bei jedem Lauf neue Zeitstempel und Dokument-IDs.
# Der feste Invariant-Modus hält Einzel- und Gesamt-PDFs byte-identisch.
rl_config.invariant = 1

# XLSX
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None  # type: ignore

# DOCX
try:
    from docx import Document
except ImportError:
    Document = None  # type: ignore

REPO_ROOT = Path(__file__).resolve().parent.parent
TESTAKTEN = REPO_ROOT / "testakten"
HELPER_DIRS = {"megaprompts"}

# Design
TEAL = HexColor("#01696F")
MUTED = HexColor("#7A7974")
BORDER = HexColor("#D4D1CA")
SURFACE = HexColor("#F7F6F2")

# Font: bevorzugt Arial, weil das Paragraph-Zeichen und deutsche Umlaute
# zuverlaessig gerendert werden. Helvetica bleibt der offline-sichere Fallback.
FONT_REG = "Helvetica"
FONT_BOLD = "Helvetica-Bold"


def _register_fonts() -> tuple[str, str]:
    candidates = [
        (
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        ),
        (
            "/Library/Fonts/Arial Unicode.ttf",
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        ),
    ]
    for regular, bold in candidates:
        regular_path = Path(regular)
        bold_path = Path(bold)
        if regular_path.exists() and bold_path.exists():
            try:
                pdfmetrics.registerFont(TTFont("AkteArial", str(regular_path)))
                pdfmetrics.registerFont(TTFont("AkteArial-Bold", str(bold_path)))
                return "AkteArial", "AkteArial-Bold"
            except Exception:
                continue
    return "Helvetica", "Helvetica-Bold"


FONT_REG, FONT_BOLD = _register_fonts()

styles = getSampleStyleSheet()
s_cover_label = ParagraphStyle(
    "CoverLabel",
    fontName=FONT_REG, fontSize=14, leading=18,
    textColor=MUTED, spaceAfter=6,
)
s_cover_title = ParagraphStyle(
    "CoverTitle",
    fontName=FONT_BOLD, fontSize=28, leading=34,
    textColor=TEAL, alignment=TA_LEFT, spaceAfter=14,
)
s_cover_sub = ParagraphStyle(
    "CoverSub",
    fontName=FONT_REG, fontSize=12, leading=16,
    textColor=black, spaceAfter=4,
)
s_cover_meta = ParagraphStyle(
    "CoverMeta",
    fontName=FONT_REG, fontSize=9, leading=12,
    textColor=MUTED, spaceAfter=3,
)
s_h1 = ParagraphStyle(
    "H1", parent=styles["Heading1"],
    fontName=FONT_BOLD, fontSize=16, leading=20, textColor=TEAL,
    spaceBefore=18, spaceAfter=8, keepWithNext=1,
)
s_h2 = ParagraphStyle(
    "H2", parent=styles["Heading2"],
    fontName=FONT_BOLD, fontSize=14, leading=18, textColor=black,
    spaceBefore=12, spaceAfter=6, keepWithNext=1,
)
s_h3 = ParagraphStyle(
    "H3", parent=styles["Heading3"],
    fontName=FONT_BOLD, fontSize=11, leading=14, textColor=black,
    spaceBefore=8, spaceAfter=4, keepWithNext=1,
)
s_body = ParagraphStyle(
    "Body", parent=styles["BodyText"],
    fontName=FONT_REG, fontSize=9.4, leading=13.2, textColor=black,
    spaceAfter=6,
)
s_meta = ParagraphStyle(
    "Meta", parent=styles["BodyText"],
    fontName=FONT_REG, fontSize=9, leading=12, textColor=MUTED,
    spaceAfter=4,
)
s_partlabel = ParagraphStyle(
    "PartLabel", parent=styles["BodyText"],
    fontName=FONT_BOLD, fontSize=11, leading=14, textColor=MUTED,
    spaceAfter=2,
)
s_doc_title = ParagraphStyle(
    "DocTitle", parent=styles["BodyText"],
    fontName=FONT_BOLD, fontSize=13, leading=16, textColor=TEAL,
    spaceAfter=4,
)
s_doc_meta = ParagraphStyle(
    "DocMeta", parent=styles["BodyText"],
    fontName=FONT_REG, fontSize=8.2, leading=10.6, textColor=black,
    spaceAfter=2,
)
s_address = ParagraphStyle(
    "Address", parent=styles["BodyText"],
    fontName=FONT_REG, fontSize=9, leading=12, textColor=black,
    leftIndent=0.15 * cm, spaceAfter=8,
)
s_quote = ParagraphStyle(
    "Quote", parent=s_body,
    leftIndent=0.55 * cm, rightIndent=0.35 * cm,
    borderColor=BORDER, borderWidth=0, borderPadding=6,
    backColor=SURFACE,
)
s_table = ParagraphStyle(
    "TableCell", parent=s_meta,
    fontName=FONT_REG, fontSize=8, leading=10,
    splitLongWords=False,
)
s_table_header = ParagraphStyle(
    "TableHeader", parent=s_table,
    fontName=FONT_BOLD,
)

SPARSE_LAST_PAGE_CHARS = 400
MAX_RENDER_ROWS = 750
MAX_RENDER_SHEETS = 40
MAX_RENDER_XML_CHARS = 300_000
MAX_RENDER_DOCX_PARAGRAPHS = 2_500
MAX_RENDER_PDF_PAGES = 1_000
COMPACT_STYLE_VALUES = {
    s_h1: {"fontSize": 15.5, "leading": 19, "spaceBefore": 14, "spaceAfter": 7},
    s_h2: {"fontSize": 13.5, "leading": 17, "spaceBefore": 10, "spaceAfter": 5},
    s_h3: {"fontSize": 10.5, "leading": 13, "spaceBefore": 6, "spaceAfter": 3},
    s_body: {"fontSize": 9.1, "leading": 12.4, "spaceAfter": 4.5},
    s_meta: {"fontSize": 8.8, "leading": 11.5, "spaceAfter": 3},
    s_table: {"fontSize": 7.8, "leading": 9.4},
    s_table_header: {"fontSize": 7.8, "leading": 9.4},
}
VERY_COMPACT_STYLE_VALUES = {
    s_h1: {"fontSize": 15, "leading": 18, "spaceBefore": 10, "spaceAfter": 6},
    s_h2: {"fontSize": 13, "leading": 16, "spaceBefore": 8, "spaceAfter": 4},
    s_h3: {"fontSize": 10.2, "leading": 12.4, "spaceBefore": 5, "spaceAfter": 2.5},
    s_body: {"fontSize": 8.8, "leading": 11.7, "spaceAfter": 3.5},
    s_meta: {"fontSize": 8.5, "leading": 10.8, "spaceAfter": 2.5},
    s_doc_title: {"fontSize": 12.2, "leading": 14.5, "spaceAfter": 3},
    s_doc_meta: {"fontSize": 8, "leading": 10, "spaceAfter": 1.5},
    s_address: {"fontSize": 8.7, "leading": 11, "spaceAfter": 6},
    s_table: {"fontSize": 7.5, "leading": 9},
    s_table_header: {"fontSize": 7.5, "leading": 9},
}


@contextmanager
def temporary_style_values(
    style_values: dict[ParagraphStyle, dict[str, object]],
):
    """Setzt globale ReportLab-Stile nur für einen einzelnen Build-Lauf."""
    original_values: list[tuple[ParagraphStyle, dict[str, object]]] = []
    for style, values in style_values.items():
        original_values.append(
            (style, {attribute: getattr(style, attribute) for attribute in values})
        )
        for attribute, value in values.items():
            setattr(style, attribute, value)
    try:
        yield
    finally:
        for style, values in original_values:
            for attribute, value in values.items():
                setattr(style, attribute, value)

# Reihenfolge der Datei-Typen im Gesamt-PDF
TYPE_ORDER = ["md", "txt", "eml", "csv", "xml", "xlsx", "docx", "image", "pdf"]
IMAGE_EXTS = {"jpg", "jpeg", "png"}
XML_EXTS = {"xml", "x83", "x84", "d83", "d84", "g48"}
TYPE_LABEL = {
    "md": "Aktenstücke",
    "txt": "Notizen und Textdateien",
    "eml": "E-Mails",
    "csv": "CSV-Tabellen",
    "xml": "XML- und GAEB-Datenanlagen",
    "xlsx": "Excel-Tabellen",
    "docx": "Word-Dokumente",
    "image": "Bildanlagen und Screenshots",
    "pdf": "PDF-Anhänge (Originaldokumente)",
}


def escape(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def md_to_flowables(md_text: str) -> list:
    out: list = []
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            out.append(Paragraph(escape(line[2:].strip()), s_h1))
            i += 1
            continue
        if line.startswith("## "):
            out.append(Paragraph(escape(line[3:].strip()), s_h2))
            i += 1
            continue
        if line.startswith("### "):
            out.append(Paragraph(escape(line[4:].strip()), s_h3))
            i += 1
            continue
        if line.startswith("#### "):
            out.append(Paragraph(escape(line[5:].strip()), s_h3))
            i += 1
            continue
        if line.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].startswith(">"):
                quote_lines.append(lines[i].lstrip("> ").rstrip())
                i += 1
            out.append(Paragraph(_inline_markup(" ".join(quote_lines).strip()), s_quote))
            out.append(Spacer(1, 4))
            continue
        if line.startswith("---"):
            out.append(Spacer(1, 6))
            i += 1
            continue
        # Tabelle?
        if (
            line.startswith("|")
            and i + 1 < len(lines)
            and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1])
        ):
            header = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            rows = [header]
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                rows.append(cells)
                i += 1
            out.extend(_render_table(rows, header=True))
            out.append(Spacer(1, 6))
            continue
        if line.startswith("- ") or line.startswith("* "):
            text = line[2:].strip()
            out.append(Paragraph("• " + _inline_markup(text), s_body))
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            out.append(Paragraph(_inline_markup(line), s_body))
            i += 1
            continue
        # Sammle normalen Absatz bis zur nächsten Leerzeile/Sondersyntax
        block = [line]
        j = i + 1
        while (
            j < len(lines)
            and lines[j].strip()
            and not lines[j].startswith(("#", "-", "*", "|", "---", ">"))
            and not re.match(r"^\d+\.\s", lines[j])
        ):
            block.append(lines[j].rstrip())
            j += 1
        text = " ".join(block).strip()
        out.append(Paragraph(_inline_markup(text), s_body))
        i = j
    return out


def _inline_markup(s: str) -> str:
    """Escape Text und setze einfache Markdown-Inline-Auszeichnungen um."""
    placeholders: list[str] = []

    def hold(value: str) -> str:
        placeholders.append(value)
        return f"@@INLINE{len(placeholders) - 1}@@"

    s = re.sub(
        r"`([^`]+)`",
        lambda m: hold(f"<font face='Courier'>{escape(m.group(1))}</font>"),
        s,
    )
    s = re.sub(
        r"\*\*([^*]+)\*\*",
        lambda m: hold(f"<b>{escape(m.group(1))}</b>"),
        s,
    )
    s = escape(s)
    for idx, value in enumerate(placeholders):
        s = s.replace(f"@@INLINE{idx}@@", value)
    return s


def _table_cell_text(s: str) -> str:
    return s.replace("<br>", "\n").replace("<br/>", "\n").strip()


def _table_cell_markup(s: str) -> str:
    text = re.sub(r"([_/-])", lambda match: match.group(1) + "\u200b", _table_cell_text(s))
    return _inline_markup(text)


def txt_to_flowables(text: str) -> list:
    out = []
    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        # Zeilenumbrueche im Absatz erhalten
        out.append(Paragraph(escape(para).replace("\n", "<br/>"), s_body))
    return out


def eml_to_flowables(path: Path) -> list:
    out = []
    try:
        with open(path, "rb") as f:
            msg = BytesParser(policy=policy.default).parse(f)
        headers = [
            ("Von", msg.get("From", "")),
            ("An", msg.get("To", "")),
            ("Datum", msg.get("Date", "")),
            ("Betreff", msg.get("Subject", "")),
        ]
        body_part = msg.get_body(preferencelist=("plain", "html"))
        body = body_part.get_content() if body_part else ""
    except Exception as e:
        out.append(Paragraph(f"<i>E-Mail konnte nicht gelesen werden: {escape(str(e))}</i>", s_meta))
        return out

    rows = [
        [Paragraph(label, s_meta), Paragraph(escape(value), s_meta)]
        for label, value in headers
    ]
    tbl = Table(rows, colWidths=[2.5 * cm, 13.5 * cm])
    tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), SURFACE),
                ("BOX", (0, 0), (-1, -1), 0.3, BORDER),
                ("INNERGRID", (0, 0), (-1, -1), 0.2, BORDER),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    out.append(tbl)
    out.append(Spacer(1, 6))
    out.extend(txt_to_flowables(body))
    return out


def csv_to_flowables(path: Path) -> list:
    out = []
    try:
        with open(path, encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = []
            truncated = False
            for row_number, row in enumerate(reader, start=1):
                if row_number > MAX_RENDER_ROWS:
                    truncated = True
                    break
                rows.append(row)
    except UnicodeDecodeError:
        with open(path, encoding="latin-1") as f:
            reader = csv.reader(f)
            rows = []
            truncated = False
            for row_number, row in enumerate(reader, start=1):
                if row_number > MAX_RENDER_ROWS:
                    truncated = True
                    break
                rows.append(row)
    except Exception as e:
        out.append(Paragraph(f"<i>CSV konnte nicht gelesen werden: {escape(str(e))}</i>", s_meta))
        return out

    if not rows:
        return out
    max_cols = max(len(r) for r in rows)
    rows = [r + [""] * (max_cols - len(r)) for r in rows]
    out.extend(_render_table(rows, header=True))
    if truncated:
        out.append(
            Paragraph(
                f"Darstellung auf {MAX_RENDER_ROWS} Zeilen begrenzt; "
                "die vollständige Datenanlage bleibt unverändert in der Akte.",
                s_meta,
            )
        )
    return out


def xlsx_to_flowables(path: Path) -> list:
    out = []
    if load_workbook is None:
        out.append(Paragraph("<i>openpyxl nicht installiert, XLSX-Inhalt wird als Originalanlage im ZIP belassen.</i>", s_meta))
        return out
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        out.append(Paragraph(f"<i>XLSX konnte nicht gelesen werden: {escape(str(e))}</i>", s_meta))
        return out
    try:
        sheet_names = wb.sheetnames[:MAX_RENDER_SHEETS]
        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            out.append(Paragraph(f"Tabellenblatt: {escape(sheet_name)}", s_h3))
            rows = []
            truncated = False
            for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if row_number > MAX_RENDER_ROWS:
                    truncated = True
                    break
                rows.append([_format_cell(c) for c in row])
            if not rows:
                continue
            # Leere Zeilen/Spalten am Ende abschneiden
            while rows and not any(c.strip() for c in rows[-1]):
                rows.pop()
            if not rows:
                continue
            max_cols = max(len(r) for r in rows)
            if max_cols == 0:
                continue
            # Hinten leere Spalten abschneiden
            while max_cols > 0 and all(
                (len(r) <= max_cols - 1) or (not r[max_cols - 1].strip()) for r in rows
            ):
                max_cols -= 1
            if max_cols == 0:
                continue
            rows = [r[:max_cols] + [""] * (max_cols - len(r[:max_cols])) for r in rows]
            out.extend(_render_table(rows, header=True))
            if truncated:
                out.append(
                    Paragraph(
                        f"Tabellenblatt auf {MAX_RENDER_ROWS} Zeilen begrenzt; "
                        "die vollständige Arbeitsmappe bleibt unverändert in der Akte.",
                        s_meta,
                    )
                )
            out.append(Spacer(1, 6))
        if len(wb.sheetnames) > MAX_RENDER_SHEETS:
            out.append(
                Paragraph(
                    f"Darstellung auf {MAX_RENDER_SHEETS} Tabellenblätter begrenzt; "
                    "die vollständige Arbeitsmappe bleibt unverändert in der Akte.",
                    s_meta,
                )
            )
    finally:
        wb.close()
    return out


def xml_to_flowables(path: Path) -> list:
    try:
        with path.open(encoding="utf-8", errors="replace") as stream:
            text = stream.read(MAX_RENDER_XML_CHARS + 1)
    except Exception as e:
        return [Paragraph(f"<i>Datenanlage konnte nicht gelesen werden: {escape(str(e))}</i>", s_meta)]
    truncated = len(text) > MAX_RENDER_XML_CHARS
    if truncated:
        text = text[:MAX_RENDER_XML_CHARS]
    text = text.replace("\t", "  ")
    out = [Paragraph(f"Datenanlage: {escape(human_filename(path))}", s_h3)]
    for block in _split_long_text(text, chunk=1200):
        out.append(Paragraph(f"<font face='Courier'>{escape(block).replace(chr(10), '<br/>')}</font>", s_meta))
        out.append(Spacer(1, 4))
    if truncated:
        out.append(
            Paragraph(
                f"Darstellung auf {MAX_RENDER_XML_CHARS:,} Zeichen begrenzt; "
                "die vollständige Datenanlage bleibt unverändert in der Akte.".replace(",", "."),
                s_meta,
            )
        )
    return out


def _format_cell(c) -> str:
    if c is None:
        return ""
    if isinstance(c, float):
        if c == int(c):
            return str(int(c))
        return f"{c:.4f}".rstrip("0").rstrip(".")
    return str(c)


# Maximalzeichen pro Zelle, ab denen die Tabelle nicht mehr als Table gerendert wird,
# sondern als sequentielle Absatzfolge (verhindert ReportLab-Overflow).
_MAX_CELL_CHARS = 1200


def _split_long_text(text: str, chunk: int = 800) -> list:
    """Schneidet sehr langen Text an Absatz- oder Satzgrenzen in Stücke."""
    text = text.replace("\r", "")
    if len(text) <= chunk:
        return [text]
    # Erst Absaetze probieren
    paras = [p for p in text.split("\n") if p.strip()]
    if any(len(p) > chunk for p in paras):
        # Weiter an Saetzen schneiden
        out = []
        for p in paras:
            if len(p) <= chunk:
                out.append(p)
                continue
            buf = ""
            for sent in p.replace("; ", "; |").replace(". ", ". |").split("|"):
                if len(sent) > chunk:
                    if buf:
                        out.append(buf)
                        buf = ""
                    out.extend(sent[start:start + chunk] for start in range(0, len(sent), chunk))
                    continue
                if len(buf) + len(sent) > chunk and buf:
                    out.append(buf)
                    buf = sent
                    continue
                buf += sent
            if buf:
                out.append(buf)
        return out
    return paras


def _render_table(rows: list, header: bool = False) -> list:
    """Rendert eine Tabelle. Falls Zellen zu lang werden, fällt es auf eine
    sequentielle Absatzdarstellung zurueck (Reihe fuer Reihe), damit ReportLab
    keine Overflow-Fehler wirft."""
    max_cell_len = max((len(c) for r in rows for c in r), default=0)
    max_cols_in_table = max((len(r) for r in rows), default=0)
    # Breite Tabellen werden als lesbare Datensätze statt als Wortsplit-Raster
    # dargestellt. Sechs kompakte Spalten bleiben auf A4 noch vergleichbar;
    # erst ab sieben Spalten wechselt die Darstellung sicherheitshalber.
    if max_cell_len > _MAX_CELL_CHARS or max_cols_in_table > 6:
        out = []
        header_row = rows[0] if header else None
        body_rows = rows[1:] if header else rows
        for ri, r in enumerate(body_rows):
            if header_row:
                record_rows = []
                for ci, cell in enumerate(r):
                    label = header_row[ci] if ci < len(header_row) else f"Spalte {ci+1}"
                    if not label.strip() and not cell.strip():
                        continue
                    record_rows.append(
                        [
                            Paragraph(escape(label or f"Spalte {ci+1}"), s_table_header),
                            Paragraph(_table_cell_markup(cell), s_table),
                        ]
                    )
                if not record_rows:
                    continue
                out.append(Paragraph(f"Datensatz {ri + 1}", s_h3))
                record = Table(record_rows, colWidths=[4.2 * cm, 11.8 * cm], splitByRow=1)
                record.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (0, -1), SURFACE),
                            ("BOX", (0, 0), (-1, -1), 0.3, BORDER),
                            ("INNERGRID", (0, 0), (-1, -1), 0.2, BORDER),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 4),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                            ("TOPPADDING", (0, 0), (-1, -1), 3),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                        ]
                    )
                )
                out.append(record)
            else:
                for ci, cell in enumerate(r):
                    for chunk in _split_long_text(cell):
                        out.append(Paragraph(_inline_markup(chunk), s_body))
            out.append(Spacer(1, 4))
        return out

    max_cols = max(len(r) for r in rows)
    avail_width = 16 * cm
    min_width = 1.35 * cm
    required_widths = []
    for col in range(max_cols):
        tokens = []
        for row in rows:
            value = row[col] if col < len(row) else ""
            tokens.extend(re.findall(r"[^\s/_-]+", value))
        widest = max(
            (pdfmetrics.stringWidth(token, FONT_REG, s_table.fontSize) for token in tokens),
            default=0,
        )
        header_width = 0
        if header and rows and col < len(rows[0]):
            header_width = pdfmetrics.stringWidth(rows[0][col], FONT_BOLD, s_table.fontSize)
        required_widths.append(max(min_width, widest + 10, header_width + 10))
    total_required = sum(required_widths)
    if total_required > avail_width:
        shrinkable = total_required - max_cols * min_width
        excess = total_required - avail_width
        factor = max(0, (shrinkable - excess) / shrinkable) if shrinkable else 0
        col_widths = [min_width + (width - min_width) * factor for width in required_widths]
    else:
        extra = avail_width - total_required
        weight = sum(required_widths) or 1
        col_widths = [width + extra * width / weight for width in required_widths]
    data = []
    for row_index, row in enumerate(rows):
        style = s_table_header if header and row_index == 0 else s_table
        data.append([Paragraph(_table_cell_markup(c), style) for c in row])
    tbl = Table(data, colWidths=col_widths, repeatRows=1 if header else 0, splitByRow=1)
    cmds = [
        ("BOX", (0, 0), (-1, -1), 0.3, BORDER),
        ("INNERGRID", (0, 0), (-1, -1), 0.2, BORDER),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    if header:
        cmds.insert(0, ("BACKGROUND", (0, 0), (-1, 0), SURFACE))
        cmds.insert(1, ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD))
    tbl.setStyle(TableStyle(cmds))
    return [tbl]


def docx_to_flowables(path: Path) -> list:
    out = []
    if Document is None:
        out.append(Paragraph("<i>python-docx nicht installiert, Inhalt wird übersprungen.</i>", s_meta))
        return out
    try:
        doc = Document(str(path))
    except Exception as e:
        out.append(Paragraph(f"<i>DOCX konnte nicht gelesen werden: {escape(str(e))}</i>", s_meta))
        return out
    for paragraph_number, para in enumerate(doc.paragraphs, start=1):
        if paragraph_number > MAX_RENDER_DOCX_PARAGRAPHS:
            out.append(
                Paragraph(
                    f"Darstellung auf {MAX_RENDER_DOCX_PARAGRAPHS} Absätze begrenzt; "
                    "das vollständige Word-Dokument bleibt unverändert in der Akte.",
                    s_meta,
                )
            )
            break
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else ""
        if style.startswith("Heading 1"):
            out.append(Paragraph(escape(text), s_h2))
        elif style.startswith("Heading 2"):
            out.append(Paragraph(escape(text), s_h3))
        elif style.startswith("Heading"):
            out.append(Paragraph(escape(text), s_h3))
        else:
            out.append(Paragraph(escape(text), s_body))
    for table in doc.tables:
        rows = []
        for r in table.rows:
            rows.append([c.text.strip() for c in r.cells])
        if rows:
            out.extend(_render_table(rows, header=True))
            out.append(Spacer(1, 6))
    return out


def image_to_flowables(path: Path) -> list:
    out = []
    try:
        width, height = ImageReader(str(path)).getSize()
        max_width = 16 * cm
        max_height = 22 * cm
        scale = min(max_width / width, max_height / height, 1)
        img = RLImage(str(path), width=width * scale, height=height * scale)
        out.append(img)
        out.append(Spacer(1, 4))
        out.append(Paragraph(f"Bildanlage: {escape(human_filename(path))}", s_meta))
    except Exception as e:
        out.append(Paragraph(f"<i>Bild konnte nicht gerendert werden: {escape(str(e))}</i>", s_meta))
    return out


def _fit_pdf_text(text: str, width: float, *, font: str = FONT_REG, size: float = 8) -> str:
    if pdfmetrics.stringWidth(text, font, size) <= width:
        return text
    suffix = " ..."
    candidate = text
    while candidate and pdfmetrics.stringWidth(candidate + suffix, font, size) > width:
        candidate = candidate[:-1]
    return candidate.rstrip() + suffix


def case_footer_text(testakte_name: str) -> str:
    defaults = TESTAKTE_DEFAULTS.get(testakte_name, {})
    aktenzeichen = defaults.get("aktenzeichen", testakte_name)
    vorhaben = defaults.get("vorhaben", "Vergabeakte")
    return f"Az. {aktenzeichen} | {vorhaben}"


def header_footer_factory(testakte_name: str):
    def hf(canv: canvas.Canvas, doc) -> None:
        canv.saveState()
        canv.setFont(FONT_REG, 8)
        canv.setFillColor(MUTED)
        footer = _fit_pdf_text(case_footer_text(testakte_name), 12.7 * cm)
        canv.drawString(2 * cm, 1.2 * cm, footer)
        canv.drawRightString(19 * cm, 1.2 * cm, f"Seite {doc.page}")
        canv.setStrokeColor(BORDER)
        canv.setLineWidth(0.3)
        canv.line(2 * cm, 1.6 * cm, 19 * cm, 1.6 * cm)
        canv.restoreState()

    return hf


def no_header_footer(canv: canvas.Canvas, doc) -> None:
    return None


def build_cover(name: str, _readme_summary: str | None, h1: str | None = None) -> list:
    """Erzeugt ein nüchternes Aktenvorblatt ohne Test- oder Bedienhinweise."""
    defaults = TESTAKTE_DEFAULTS.get(name, {})
    title = defaults.get("vorhaben") or h1 or name.replace("-", " ").title()
    rows = [
        [Paragraph("Aktenzeichen", s_table_header), Paragraph(escape(defaults.get("aktenzeichen", "-")), s_table)],
        [Paragraph("Auftraggeber", s_table_header), Paragraph(escape(defaults.get("auftraggeber", "-")), s_table)],
        [Paragraph("Vorhaben", s_table_header), Paragraph(escape(title), s_table)],
        [Paragraph("Band", s_table_header), Paragraph("Aktenstücke und Anlagen", s_table)],
    ]
    table = Table(rows, colWidths=[4.0 * cm, 12.0 * cm], splitByRow=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), SURFACE),
                ("BOX", (0, 0), (-1, -1), 0.4, BORDER),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, BORDER),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return [
        Spacer(1, 3.0 * cm),
        Paragraph("VERGABEAKTE", s_cover_label),
        Paragraph(escape(title), s_cover_title),
        HRFlowable(width="100%", thickness=1.2, color=TEAL, spaceAfter=18),
        table,
        Spacer(1, 9.0 * cm),
        Paragraph("Dokumentenband", s_cover_meta),
        PageBreak(),
    ]


def extract_readme_summary(readme_path: Path) -> tuple[str | None, str | None]:
    """Liest aus der README den H1-Titel und einen kurzen beschreibenden Absatz.

    Der beschreibende Absatz wird absichtlich erst gesucht, NACHDEM Download-Bloecke
    und Aktenstruktur-Blocks übersprungen wurden, damit nicht der ZIP-Hinweis
    auf dem Cover landet.
    """
    if not readme_path.is_file():
        return None, None
    text = readme_path.read_text(encoding="utf-8")
    # H1
    h1_match = re.search(r"^#\s+(.+?)\s*$", text, flags=re.MULTILINE)
    h1 = h1_match.group(1).strip() if h1_match else None

    # Suche eine Sektion mit beschreibendem Inhalt: 'Kurzbild', 'Worum',
    # 'Sachverhalt', 'Überblick', 'Mandat', 'Fall' o.ae.
    section_pattern = re.compile(
        r"^##[^\n]*?(?:kurzbild|worum geht|sachverhalt|überblick|\u00fcberblick|mandat|fall|der fall|akte|kontext|ausgangslage|ausgangs|zweck|szenario|idee|einsatz|\u00fcbersicht|uebersicht|verfahrenseckdaten|aktenkern|aktenbestand|mandantenkonstellation|politische vorgabe|enthaltene arbeitsdateien|dateien)[^\n]*\n([\s\S]*?)(?=^## |\Z)",
        re.IGNORECASE | re.MULTILINE,
    )
    m = section_pattern.search(text)
    candidate_text = m.group(1) if m else text
    for para in candidate_text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        if para.startswith(("#", "-", "*", "|", "<!--", "```")):
            continue
        # Download-/ZIP-Hinweise ueberspringen
        lower = para.lower()
        if any(
            kw in lower
            for kw in (
                "zip-datei",
                "zip datei",
                "direkt-download",
                "als zip",
                "github-release",
                "github release",
                "download",
            )
        ):
            continue
        para = re.sub(r"\*\*([^*]+)\*\*", r"\1", para)
        para = re.sub(r"\s+", " ", para)
        return h1, para[:400]
    return h1, None


DERIVED_SOURCE_DIRS = {"docx", "eml", "xlsx", "scan-pdf", "pptx"}


def collect_files(testakte_dir: Path, *, canonical_only: bool = False) -> dict[str, list[Path]]:
    files_by_type: dict[str, list[Path]] = {t: [] for t in TYPE_ORDER}
    root_source_stems = {
        path.stem
        for path in testakte_dir.glob("[0-9][0-9]_*.md")
        if path.is_file()
    }
    for f in testakte_dir.rglob("*"):
        if not include_in_working_dump(f, testakte_dir):
            continue
        rel = f.relative_to(testakte_dir)
        if (
            canonical_only
            and len(rel.parts) > 1
            and rel.parts[0] in DERIVED_SOURCE_DIRS
            and f.stem in root_source_stems
        ):
            continue
        ext = f.suffix.lower().lstrip(".")
        if ext in IMAGE_EXTS:
            files_by_type["image"].append(f)
            continue
        if ext in XML_EXTS:
            files_by_type["xml"].append(f)
            continue
        if ext not in TYPE_ORDER:
            continue
        files_by_type[ext].append(f)
    if not any(files_by_type.values()):
        readme = testakte_dir / "README.md"
        if readme.is_file():
            files_by_type["md"].append(readme)
    for t in files_by_type:
        files_by_type[t].sort(key=lambda p: str(p.relative_to(testakte_dir)).lower())
    return files_by_type


def raw_text_for_meta(path: Path, t: str) -> str:
    if t == "md":
        return path.read_text(encoding="utf-8", errors="replace")
    if t in {"txt", "csv", "eml", "xml"}:
        return path.read_text(encoding="utf-8", errors="replace")
    return path.name


HUMAN_FILENAME_REPLACEMENTS = (
    ("nachpruefung", "Nachprüfung"),
    ("pruefvermerk", "Prüfvermerk"),
    ("pruefmatrix", "Prüfmatrix"),
    ("aufklaerung", "Aufklärung"),
    ("geschaeftsgeheimnis", "Geschäftsgeheimnis"),
    ("veroeffentlichung", "Veröffentlichung"),
    ("uebertragung", "Übertragung"),
    ("uebergabe", "Übergabe"),
    ("rueckversetzung", "Rückversetzung"),
    ("eignungspruefung", "Eignungsprüfung"),
    ("aktivitaetsprotokoll", "Aktivitätsprotokoll"),
    ("auswaertung", "Auswertung"),
    ("aenderung", "Änderung"),
    ("flaechen", "Flächen"),
    ("ruege", "Rüge"),
)

FRIENDLY_SOURCE_TITLES = {
    "00_bearbeitungsnotizen_eingangskorb": "Bearbeitungsnotizen zum Akteneingang",
    "00_interne_weiterleitung_aktenstand": "Interne Weiterleitung zum Aktenstand",
    "00_portal_aktivitaetsprotokoll": "Portal-Aktivitätsprotokoll",
    "00_fristen_upload_skizze": "Fristen- und Uploadübersicht",
    "reinigungsflaechen_frequenzen": "Reinigungsflächen und Leistungsfrequenzen",
    "eforms_contract_notice_reinigung": "eForms-Auftragsbekanntmachung Reinigung",
    "preisblatt_laptops": "Preisblatt mobile Endgeräte",
    "eforms_contract_notice_laptops": "eForms-Auftragsbekanntmachung mobile Endgeräte",
    "fristenmatrix": "Fristenmatrix",
    "angebot_hm_x84": "GAEB-Angebot Hochbau Mustermann X84",
    "lv_stadthalle_x83": "GAEB-Leistungsverzeichnis Stadthalle X83",
    "auswahlmatrix": "Auswahlmatrix Teilnahmewettbewerb",
    "teilnahmeantrag_bietergemeinschaft": "Teilnahmeantrag der Bietergemeinschaft",
    "09_pruefmatrix_auftragnehmerwechsel": "Prüfmatrix zum Auftragnehmerwechsel",
    "10_vk_fristen_und_kosten": "Fristen- und Kostenmatrix der Nachprüfung",
    "angriffslinien_belegmatrix": "Angriffslinien- und Belegmatrix",
    "fristenampel": "Fristenampel",
    "wertungsmatrix_auszug": "Auszug aus der Wertungsmatrix",
    "eforms_change_notice": "eForms-Änderungsbekanntmachung",
    "dokumentenmatrix_upload_export": "Dokumentenmatrix für den Upload",
    "streitfall_dashboard": "Streitfall-Dashboard",
    "bewertungsmatrix_zuschlagskriterien": "Bewertungsmatrix der Zuschlagskriterien",
    "eignungsreferenzen_cybershield": "Eignungsreferenzen CyberShield Defense GmbH",
    "bieter_konstellation": "Bieterkonstellation",
    "bsi_schutzziel_pyramide": "BSI-Schutzziele",
    "organigramm_vergabestelle": "Organigramm der Vergabestelle",
    "ted_bekanntmachung_auszug": "TED-Auftragsbekanntmachung",
    "vergabevermerk_auszug": "Auszug aus dem Vergabevermerk",
}


def human_filename(path: Path) -> str:
    if path.stem in FRIENDLY_SOURCE_TITLES:
        return FRIENDLY_SOURCE_TITLES[path.stem]
    stem = re.sub(r"^\d{2}_", "", path.stem.casefold())
    for source, target in HUMAN_FILENAME_REPLACEMENTS:
        stem = stem.replace(source, target)
    words = re.split(r"[_-]+", stem)
    acronyms = {
        "bsi": "BSI",
        "eforms": "eForms",
        "gaeb": "GAEB",
        "gwb": "GWB",
        "it": "IT",
        "olg": "OLG",
        "soc": "SOC",
        "ted": "TED",
        "vk": "VK",
        "x83": "X83",
        "x84": "X84",
        "d83": "D83",
        "d84": "D84",
        "g48": "G48",
    }
    rendered = [acronyms.get(word, word) for word in words if word]
    title = " ".join(rendered).strip()
    if title:
        title = title[0].upper() + title[1:]
    return title or "Aktenstück"


def source_properties(path: Path) -> dict[str, str]:
    props: dict[str, str] = {}
    suffix = path.suffix.lower()
    if suffix == ".eml":
        try:
            msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            props["betreff"] = str(msg.get("Subject", "")).strip()
            props["absender"] = str(msg.get("From", "")).strip()
            props["empfaenger"] = str(msg.get("To", "")).strip()
            props["aktenzeichen"] = str(msg.get("X-Aktenzeichen", "")).strip()
            sent = parsedate_to_datetime(str(msg.get("Date", "")))
            props["datum"] = sent.strftime("%d.%m.%Y")
        except Exception:
            pass
    elif suffix == ".txt":
        try:
            props["betreff"] = next(
                (
                    line.strip()
                    for line in path.read_text(encoding="utf-8").splitlines()
                    if line.strip()
                ),
                "",
            )
        except (OSError, UnicodeError):
            pass
    elif suffix == ".docx" and Document is not None:
        try:
            source = Document(str(path))
            props["betreff"] = source.core_properties.title or next(
                (
                    paragraph.text.strip()
                    for paragraph in source.paragraphs
                    if paragraph.text.strip()
                ),
                "",
            )
        except Exception:
            pass
    elif suffix == ".xlsx" and load_workbook is not None:
        try:
            workbook = load_workbook(path, read_only=True)
            props["betreff"] = workbook.properties.title or ""
            workbook.close()
        except Exception:
            pass
    elif suffix == ".pdf":
        try:
            props["betreff"] = (PdfReader(str(path)).metadata.title or "").strip()
        except Exception:
            pass
    if props.get("betreff") in {path.name, path.stem}:
        props["betreff"] = ""
    props["betreff"] = props.get("betreff") or human_filename(path)
    return props


def render_body_flowables(path: Path, t: str) -> list:
    if t == "md":
        return md_to_flowables(markdown_text_for_pdf(path))
    if t == "txt":
        return txt_to_flowables(path.read_text(encoding="utf-8", errors="replace"))
    if t == "eml":
        return eml_to_flowables(path)
    if t == "csv":
        return csv_to_flowables(path)
    if t == "xml":
        return xml_to_flowables(path)
    if t == "xlsx":
        return xlsx_to_flowables(path)
    if t == "docx":
        return docx_to_flowables(path)
    if t == "image":
        return image_to_flowables(path)
    return []


def keep_short_tail_together(
    flowables: list,
    *,
    max_items: int = 12,
    target_chars: int = 700,
) -> list:
    """Hält einen kurzen Schlussblock zusammen, ohne Tabellen zu blockieren."""
    start = len(flowables)
    paragraph_count = 0
    text_chars = 0
    while (
        start > 0
        and len(flowables) - start < max_items
        and text_chars < target_chars
    ):
        candidate = flowables[start - 1]
        if not isinstance(candidate, (Paragraph, Spacer, HRFlowable)):
            break
        start -= 1
        if isinstance(candidate, Paragraph):
            paragraph_count += 1
            text_chars += len(re.sub(r"\s+", "", candidate.getPlainText()))
    if paragraph_count < 2 or start == len(flowables):
        return flowables
    return [*flowables[:start], KeepTogether(flowables[start:])]


def single_pdf_name(testakte_dir: Path, path: Path) -> str:
    rel = path.relative_to(testakte_dir)
    base = rel.with_suffix("").as_posix().replace("/", "__")
    base = re.sub(r"[^A-Za-z0-9_.-]+", "-", base)
    return f"{base}.pdf"


def pdf_subject(meta: dict[str, str]) -> str:
    return f"{meta['dokumenttyp']} | Az. {meta['aktenzeichen']}"


def _build_text_individual_pdf(
    testakte_dir: Path,
    path: Path,
    t: str,
    out_path: Path,
    raw_text: str,
    *,
    compact: int,
) -> None:
    style_values = (
        VERY_COMPACT_STYLE_VALUES
        if compact >= 2
        else COMPACT_STYLE_VALUES
        if compact
        else {}
    )
    with temporary_style_values(style_values):
        margin_top = (1.3 if compact >= 2 else 1.5 if compact else 1.8) * cm
        margin_bottom = (1.4 if compact >= 2 else 1.55 if compact else 1.9) * cm
        meta = document_meta(testakte_dir, path, raw_text)
        doc = SimpleDocTemplate(
            str(out_path),
            pagesize=A4,
            leftMargin=2.15 * cm,
            rightMargin=2.15 * cm,
            topMargin=margin_top,
            bottomMargin=margin_bottom,
            title=meta["betreff"],
            author=meta["absender"],
            subject=pdf_subject(meta),
            keywords=f"Vergabeakte, {meta['aktenzeichen']}, {meta['dokumenttyp']}",
        )
        flow: list = []
        flow.extend(document_header_flowables(testakte_dir, path, raw_text))
        flow.extend(render_body_flowables(path, t))
        hf = header_footer_factory(testakte_dir.name)
        doc.build(flow, onFirstPage=hf, onLaterPages=hf)


def _sparse_last_page(path: Path) -> bool:
    reader = PdfReader(str(path))
    if len(reader.pages) <= 1:
        return False
    text = reader.pages[-1].extract_text() or ""
    return len(re.sub(r"\s+", "", text)) < SPARSE_LAST_PAGE_CHARS


def build_individual_source_pdf(testakte_dir: Path, path: Path, t: str, out_dir: Path) -> tuple[bool, str]:
    out_path = out_dir / single_pdf_name(testakte_dir, path)
    raw_text = raw_text_for_meta(path, t)
    meta = document_meta(testakte_dir, path, raw_text)
    if t == "pdf":
        writer = PdfWriter()
        append_pdf_with_separator(
            writer,
            f"Originalanlage: {meta['betreff']}",
            path,
            testakte_dir.name,
        )
        writer.add_metadata(
            {
                "/Title": meta["betreff"],
                "/Author": meta["absender"],
                "/Subject": pdf_subject(meta),
                "/Keywords": f"Vergabeakte, {meta['aktenzeichen']}, {meta['dokumenttyp']}",
            }
        )
        with open(out_path, "wb") as f:
            writer.write(f)
        return True, out_path.name

    try:
        _build_text_individual_pdf(
            testakte_dir,
            path,
            t,
            out_path,
            raw_text,
            compact=0,
        )
        if _sparse_last_page(out_path):
            _build_text_individual_pdf(
                testakte_dir,
                path,
                t,
                out_path,
                raw_text,
                compact=1,
            )
        if _sparse_last_page(out_path):
            _build_text_individual_pdf(
                testakte_dir,
                path,
                t,
                out_path,
                raw_text,
                compact=2,
            )
        return True, out_path.name
    except Exception as e:
        return False, f"{path.relative_to(testakte_dir)}: {e}"


def build_individual_pdfs(testakte_dir: Path, files: dict[str, list[Path]]) -> tuple[int, list[str]]:
    if testakte_dir.name in {"megaprompts"}:
        return 0, []
    out_dir = testakte_dir / "einzel-pdf"
    staging_dir = Path(
        tempfile.mkdtemp(prefix=".einzel-pdf-", dir=testakte_dir)
    )
    errors: list[str] = []
    count = 0
    try:
        for t in TYPE_ORDER:
            for path in files[t]:
                ok, info = build_individual_source_pdf(testakte_dir, path, t, staging_dir)
                if ok:
                    count += 1
                else:
                    errors.append(info)
        if errors:
            return count, errors
        backup_dir = testakte_dir / ".einzel-pdf-backup"
        if backup_dir.exists():
            shutil.rmtree(backup_dir)
        if out_dir.exists():
            out_dir.replace(backup_dir)
        try:
            staging_dir.replace(out_dir)
        except Exception:
            if backup_dir.exists() and not out_dir.exists():
                backup_dir.replace(out_dir)
            raise
        if backup_dir.exists():
            shutil.rmtree(backup_dir)
    finally:
        if staging_dir.exists():
            shutil.rmtree(staging_dir)
    return count, errors


AUTOGEN_BLOCK_RE = re.compile(
    r"<!-- BEGIN [^\n]*?-->\s*[\s\S]*?<!-- END [^\n]*?-->\s*",
    re.MULTILINE,
)
AKTENMETA_RE = re.compile(r"<!--\s*aktenmeta\s*([\s\S]*?)-->", re.IGNORECASE)

TESTAKTE_DEFAULTS = {
    "01-reinigung-grundschule-musterstadt": {
        "auftraggeber": "Stadt Musterstadt - Zentrale Vergabestelle",
        "auftraggeber_anschrift": "Rathausplatz 1, 99999 Musterstadt",
        "aktenzeichen": "VS-2026-RE-014",
        "vorhaben": "Unterhalts- und Glasreinigung Grundschulen",
    },
    "02-laptopbeschaffung-landesamt-musterland": {
        "auftraggeber": "Landesamt für Digitalisierung Musterland",
        "auftraggeber_anschrift": "Dienstgebäude Nord, Musterallee 10, 99000 Musterstadt",
        "aktenzeichen": "LAD-IT-2026-044",
        "vorhaben": "Rahmenvertrag mobile Endgeräte 2026",
    },
    "03-bauleistung-stadthalle-musterstadt": {
        "auftraggeber": "Stadt Musterstadt - Fachbereich Hochbau",
        "auftraggeber_anschrift": "Bauhofstraße 4, 99999 Musterstadt",
        "aktenzeichen": "HB-STH-2026-019",
        "vorhaben": "Sanierung Stadthalle Musterstadt, Dach und Gebäudehülle",
    },
    "04-it-beratung-ministerium-musterland": {
        "auftraggeber": "Ministerium für Inneres und Digitalisierung Musterland",
        "auftraggeber_anschrift": "Regierungsufer 3, 99000 Musterstadt",
        "aktenzeichen": "MID-STRAT-2026-07",
        "vorhaben": "Cybersicherheitsberatung mit 24/7-Reaktionsdienst",
    },
    "05-insolvenz-auftragnehmerwechsel-132-gwb": {
        "auftraggeber": "Stadt Hafenburg - Amt für Digitalisierung",
        "auftraggeber_anschrift": "Markt 7, 21000 Hafenburg",
        "aktenzeichen": "HB-IT-SERVICE-2026",
        "vorhaben": "IT-Betrieb kommunaler Fachverfahren nach Auftragnehmerinsolvenz",
    },
    "06-konkurrentenrechtsschutz-rechenzentrum-musterkreis": {
        "auftraggeber": "Musterkreis IT-Service",
        "auftraggeber_anschrift": "Kreishaus 1, 44100 Musterstadt",
        "aktenzeichen": "DW-2026-RZ-011",
        "vorhaben": "Rechenzentrumsbetrieb und Konkurrentenrechtsschutz",
    },
    "it-sig-2-vergabe-landeshauptstadt-schwerin-nachpruefung": {
        "auftraggeber": "Landeshauptstadt Schwerin - Eigenbetrieb Stadtwirtschaft Schwerin SDS",
        "auftraggeber_anschrift": "Am Grünen Tal 18, 19063 Schwerin",
        "aktenzeichen": "LH-SN-Cyber-SOC-NSV-2026",
        "vorhaben": "SOC Managed Service und ISB-Beratung",
    },
}


def parse_aktenmeta(text: str) -> dict[str, str]:
    match = AKTENMETA_RE.search(text)
    if not match:
        return {}
    meta: dict[str, str] = {}
    for line in match.group(1).splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = key.strip().lower()
        key = (
            key.replace("ä", "ae")
            .replace("ö", "oe")
            .replace("ü", "ue")
            .replace("ß", "ss")
        )
        value = value.strip()
        if key and value:
            meta[key] = value
    return meta


def strip_aktenmeta(text: str) -> str:
    return AKTENMETA_RE.sub("", text)


def _first_h1(text: str) -> str | None:
    m = re.search(r"^#\s+(.+?)\s*$", text, flags=re.MULTILINE)
    return m.group(1).strip() if m else None


def _first_date(text: str) -> str:
    m = re.search(r"\b\d{2}\.\d{2}\.\d{4}\b", text)
    if m:
        return m.group(0)
    iso = re.search(r"\b(20\d{2})-(\d{2})-(\d{2})\b", text)
    if iso:
        return f"{iso.group(3)}.{iso.group(2)}.{iso.group(1)}"
    return "-"


def _first_aktenzeichen(text: str, fallback: str) -> str:
    patterns = [
        r"Aktenzeichen\s*[:\-]\s*([A-ZÄÖÜa-zäöü0-9./\- ]{4,60})",
        r"Geschäftszeichen\s*[:\-]\s*([A-ZÄÖÜa-zäöü0-9./\- ]{4,60})",
        r"Vergabenummer\s*[:\-]\s*([A-ZÄÖÜa-zäöü0-9./\- ]{4,60})",
    ]
    for pattern in patterns:
        m = re.search(pattern, text)
        if m:
            return m.group(1).strip().rstrip(".")
    return fallback


def infer_document_type(path: Path, text: str) -> str:
    stem = path.stem.lower()
    title = (_first_h1(text) or "").lower()
    hay = f"{stem} {title}"
    if "fristen_und_kosten" in stem:
        return "Fristen- und Kostenmatrix"
    if "pruefmatrix" in stem or "prüfmatrix" in stem:
        return "Prüfmatrix"
    rules = [
        ("bekanntmachung", "Auftragsbekanntmachung / eForms-Datenblatt"),
        ("bedarfs", "Bedarfsvermerk"),
        ("auftragswert", "Auftragswertschätzung"),
        ("leistungsbeschreibung", "Leistungsbeschreibung / Leistungsverzeichnis"),
        ("eignung", "Eignungs- und Zuschlagskriterien"),
        ("angebot", "Angebotsunterlage"),
        ("referenz", "Referenznachweis"),
        ("ausschluss", "Ausschlussschreiben"),
        ("ruege", "Rügeschreiben"),
        ("rüge", "Rügeschreiben"),
        ("nichtabhilfe", "Nichtabhilfeentscheidung"),
        ("nachpruefung", "Nachprüfungsantrag"),
        ("nachprüfung", "Nachprüfungsantrag"),
        ("stellungnahme", "Stellungnahme"),
        ("erwiderung", "Erwiderung"),
        ("beschluss", "Beschluss"),
        ("aufklaerung", "Aufklärungsverfügung"),
        ("aufklärung", "Aufklärungsverfügung"),
        ("termin", "Terminsprotokoll"),
        ("vergleich", "Vergleichsvorschlag"),
        ("beschwerde", "Sofortige Beschwerde"),
        ("klage", "Antragsschrift"),
        ("kosten", "Kostenvermerk"),
        ("sachverhalt", "Sachverhalt und Chronologie"),
        ("pruefvermerk", "Prüfvermerk"),
        ("prüfvermerk", "Prüfvermerk"),
        ("wertung", "Wertungsvermerk"),
        ("bieterfrage", "Bieterfrage und Antwortentwurf"),
        ("beilad", "Beiladung / Beiladungsantrag"),
        ("eml", "E-Mail"),
    ]
    for needle, label in rules:
        if needle in hay:
            return label
    if path.parent.name == "notizen":
        return "Interne Bearbeitungsnotiz"
    if "portal_aktivitaetsprotokoll" in stem:
        return "Portal-Aktivitätsprotokoll"
    if path.suffix.lower() == ".csv":
        return "CSV-Datenanlage"
    if path.suffix.lower() == ".txt":
        return "Textnotiz"
    if path.suffix.lower() == ".eml":
        return "E-Mail"
    if path.suffix.lower() == ".xlsx":
        return "Tabellenanlage"
    if path.suffix.lower().lstrip(".") in XML_EXTS:
        return "XML-/GAEB-Datenanlage"
    if path.suffix.lower() == ".docx":
        return "Schriftsatzanlage"
    if path.suffix.lower() in {".jpg", ".jpeg", ".png"}:
        return "Bildanlage"
    if path.suffix.lower() == ".pdf":
        return "PDF-Originalanlage"
    return "Aktenstück"


def infer_sender(path: Path, text: str, defaults: dict[str, str]) -> str:
    hay = f"{path.stem.lower()} {text[:1500].lower()}"
    if path.parent.name in {"notizen", "bilder"}:
        return "Interne Bearbeitung"
    if path.suffix.lower() in {".csv", ".xlsx"} and any(
        marker in path.stem.lower() for marker in ("pruefmatrix", "fristen_und_kosten")
    ):
        return defaults.get("auftraggeber", "Vergabestelle")
    if path.suffix.lower().lstrip(".") in XML_EXTS:
        if any(token in path.stem.lower() for token in ("angebot", "x84", "d84")):
            return "Bieterseite / Verfahrensbevollmächtigte"
        return defaults.get("auftraggeber", "Vergabestelle")
    if any(x in hay for x in ("vergabekammer", "beschluss vk", "vk bund")):
        return "Vergabekammer / Gericht"
    if any(x in hay for x in ("rüge", "ruege", "nachprüfungsantrag", "nachpruefungsantrag", "beschwerde", "klageschrift", "cybershield", "dataportus", "bialto")):
        return "Bieterseite / Verfahrensbevollmächtigte"
    return defaults.get("auftraggeber", "Vergabestelle")


def infer_recipient(path: Path, text: str, defaults: dict[str, str]) -> str:
    hay = f"{path.stem.lower()} {text[:1500].lower()}"
    if path.parent.name in {"notizen", "bilder"}:
        return "Interne Vergabeakte"
    if path.suffix.lower() in {".csv", ".xlsx"} and any(
        marker in path.stem.lower() for marker in ("pruefmatrix", "fristen_und_kosten")
    ):
        return "Interne Vergabeakte"
    if any(x in hay for x in ("nachprüfungsantrag", "nachpruefungsantrag", "beschwerde", "klageschrift")):
        return "Vergabekammer bzw. zuständiges Beschwerdegericht"
    if any(x in hay for x in ("rüge", "ruege", "angebot", "referenz")):
        return defaults.get("auftraggeber", "Vergabestelle")
    if any(x in hay for x in ("ausschluss", "nichtabhilfe", "aufklärung", "aufklaerung")):
        return "Bieter / Bewerber laut Verteiler"
    return "Vergabeakte"


def document_meta(testakte_dir: Path, path: Path, text: str) -> dict[str, str]:
    defaults = TESTAKTE_DEFAULTS.get(testakte_dir.name, {})
    parsed = parse_aktenmeta(text)
    source = source_properties(path)
    document_type = (
        parsed.get("dokumenttyp")
        or source.get("dokumenttyp")
        or infer_document_type(path, text)
    )
    title = parsed.get("betreff") or _first_h1(text) or source.get("betreff") or human_filename(path)
    if title == path.name and document_type in {"Fristen- und Kostenmatrix", "Prüfmatrix"}:
        title = document_type
    fallback_az = defaults.get("aktenzeichen", testakte_dir.name)
    meta = {
        "dokumenttyp": document_type,
        "absender": (
            parsed.get("absender")
            or source.get("absender")
            or infer_sender(path, text, defaults)
        ),
        "empfaenger": (
            parsed.get("empfaenger")
            or source.get("empfaenger")
            or infer_recipient(path, text, defaults)
        ),
        "datum": parsed.get("datum") or source.get("datum") or _first_date(text),
        "aktenzeichen": (
            parsed.get("aktenzeichen")
            or source.get("aktenzeichen")
            or _first_aktenzeichen(text, fallback_az)
        ),
        "betreff": title,
        "vorhaben": parsed.get("vorhaben") or defaults.get("vorhaben", testakte_dir.name),
    }
    return meta


def document_header_flowables(testakte_dir: Path, path: Path, text: str) -> list:
    meta = document_meta(testakte_dir, path, text)
    defaults = TESTAKTE_DEFAULTS.get(testakte_dir.name, {})
    sender = meta["absender"]
    if sender == defaults.get("auftraggeber") and defaults.get("auftraggeber_anschrift"):
        sender = f"{sender}<br/>{defaults['auftraggeber_anschrift']}"
    number = re.match(r"^(\d{2})_", path.stem) if path.parent == testakte_dir else None
    reference = (
        f"{number.group(1)} | {meta['dokumenttyp']}"
        if number
        else human_filename(path)
    )
    right_lines = [
        f"<b>{escape(meta['dokumenttyp'])}</b>",
        f"Aktenzeichen: {escape(meta['aktenzeichen'])}",
    ]
    if meta["datum"] != "-":
        right_lines.append(f"Datum: {escape(meta['datum'])}")
    right_lines.append(f"Aktenstück: {escape(reference)}")
    right = "<br/>".join(right_lines)
    top = Table(
        [
            [
                Paragraph(sender.replace("\n", "<br/>"), s_doc_meta),
                Paragraph(right, s_doc_meta),
            ]
        ],
        colWidths=[9.4 * cm, 6.6 * cm],
    )
    top.setStyle(
        TableStyle(
            [
                ("BOX", (1, 0), (1, 0), 0.35, BORDER),
                ("BACKGROUND", (1, 0), (1, 0), SURFACE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return [
        top,
        Spacer(1, 9),
        Paragraph(escape(meta["empfaenger"]).replace("\n", "<br/>"), s_address),
        Paragraph(escape(meta["betreff"]), s_doc_title),
        Paragraph(f"Vorhaben: {escape(meta['vorhaben'])}", s_meta),
        HRFlowable(width="100%", thickness=0.45, color=BORDER),
        Spacer(1, 9),
    ]


def markdown_text_for_pdf(path: Path) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.name == "README.md":
        text = AUTOGEN_BLOCK_RE.sub("", text)
    text = strip_aktenmeta(text)
    return text.strip()


class PageStartMarker(Flowable):
    def __init__(self, record: dict[str, int | None]):
        super().__init__()
        self.width = 0
        self.height = 0
        self.record = record

    def draw(self) -> None:
        self.record["page"] = max(0, self.canv.getPageNumber() - 1)


def outline_label(path: Path, meta: dict[str, str]) -> str:
    number = re.match(r"^(\d{2})_", path.stem)
    prefix = f"{number.group(1)} | " if number else ""
    label = prefix + meta["betreff"]
    return _fit_pdf_text(label, 16 * cm, size=9)


def build_text_pdf(
    testakte_dir: Path,
    files: dict[str, list[Path]],
    cover: list,
    tmp_path: Path,
) -> tuple[bool, list[Path], list[tuple[str, int]]]:
    """Baut den Text-Teil als PDF, sammelt PDF-Anhaenge separat."""
    defaults = TESTAKTE_DEFAULTS.get(testakte_dir.name, {})
    doc = SimpleDocTemplate(
        str(tmp_path),
        pagesize=A4,
        leftMargin=2.15 * cm, rightMargin=2.15 * cm,
        topMargin=1.3 * cm, bottomMargin=1.4 * cm,
        title=defaults.get("vorhaben", testakte_dir.name),
        author=defaults.get("auftraggeber", "Vergabeakte"),
        subject=f"Vergabeakte | Az. {defaults.get('aktenzeichen', testakte_dir.name)}",
        keywords=f"Vergabeakte, {defaults.get('aktenzeichen', testakte_dir.name)}",
    )
    flow = list(cover)

    pdf_attachments: list[Path] = []
    bookmark_records: list[tuple[str, dict[str, int | None]]] = []
    for t in TYPE_ORDER:
        if not files[t]:
            continue
        if t == "pdf":
            # PDFs werden separat angehaengt (Original-Layout bewahren)
            pdf_attachments = files[t]
            continue
        for f in files[t]:
            try:
                raw_text = raw_text_for_meta(f, t)
                meta = document_meta(testakte_dir, f, raw_text)
                record: dict[str, int | None] = {"page": None}
                bookmark_records.append((outline_label(f, meta), record))
                flow.append(PageStartMarker(record))
                flow.extend(document_header_flowables(testakte_dir, f, raw_text))
                flow.extend(keep_short_tail_together(render_body_flowables(f, t)))
            except Exception as e:
                flow.append(Paragraph(f"<i>Inhalt konnte nicht gerendert werden: {escape(str(e))}</i>", s_meta))
            # Jedes Aktenstück beginnt im Gesamt-PDF auf einer neuen Seite
            flow.append(PageBreak())

    if not flow:
        flow.append(Paragraph("Dateiablage: Original-PDFs folgen.", s_meta))

    # Trailing PageBreak entfernen, damit am Ende keine Leerseite entsteht.
    while flow and isinstance(flow[-1], PageBreak):
        flow.pop()

    hf = header_footer_factory(testakte_dir.name)
    try:
        # Jedes Aktenstück beginnt auf einer neuen Seite. Die für schwierige
        # Einzelstücke bewährten Aktenstile verhindern isolierte Schlusszeilen.
        with temporary_style_values(VERY_COMPACT_STYLE_VALUES):
            doc.build(flow, onFirstPage=hf, onLaterPages=hf)
    except Exception as e:
        print(f"  FEHLER beim Bauen: {e}")
        return False, pdf_attachments, []
    bookmarks = [
        (label, int(record["page"]))
        for label, record in bookmark_records
        if record["page"] is not None
    ]
    return True, pdf_attachments, bookmarks


def append_pdf_with_separator(
    writer: PdfWriter,
    label: str,
    pdf_path: Path,
    testakte_name: str,
) -> int:
    start_page = len(writer.pages)
    sep = io.BytesIO()
    c = canvas.Canvas(sep, pagesize=A4)
    c.setTitle(label)
    c.setAuthor(TESTAKTE_DEFAULTS.get(testakte_name, {}).get("auftraggeber", "Vergabeakte"))
    c.setFont(FONT_BOLD, 14)
    c.setFillColor(TEAL)
    c.drawString(2 * cm, 25 * cm, "Originalanlage")
    c.setFont(FONT_REG, 9)
    c.setFillColor(MUTED)
    c.drawString(2 * cm, 24.2 * cm, _fit_pdf_text(label, 16.5 * cm, size=9))
    c.setStrokeColor(BORDER)
    c.setLineWidth(0.3)
    c.line(2 * cm, 1.6 * cm, 19 * cm, 1.6 * cm)
    c.setFont(FONT_REG, 8)
    c.drawString(
        2 * cm,
        1.2 * cm,
        _fit_pdf_text(case_footer_text(testakte_name), 16.5 * cm),
    )
    c.showPage()
    c.save()
    sep.seek(0)
    for p in PdfReader(sep).pages:
        writer.add_page(p)
    try:
        source_reader = PdfReader(str(pdf_path))
        for page_number, p in enumerate(source_reader.pages, start=1):
            if page_number > MAX_RENDER_PDF_PAGES:
                break
            writer.add_page(p)
        if len(source_reader.pages) > MAX_RENDER_PDF_PAGES:
            notice = io.BytesIO()
            notice_canvas = canvas.Canvas(notice, pagesize=A4)
            notice_canvas.setFont(FONT_REG, 10)
            notice_canvas.drawString(
                2 * cm,
                25 * cm,
                f"Darstellung auf {MAX_RENDER_PDF_PAGES} Seiten begrenzt.",
            )
            notice_canvas.drawString(
                2 * cm,
                24.3 * cm,
                "Die vollständige Originalanlage bleibt unverändert in der Akte.",
            )
            notice_canvas.showPage()
            notice_canvas.save()
            notice.seek(0)
            for page in PdfReader(notice).pages:
                writer.add_page(page)
    except Exception as e:
        # PDF defekt oder verschluesselt -> Hinweisseite einfügen
        sep2 = io.BytesIO()
        c2 = canvas.Canvas(sep2, pagesize=A4)
        c2.setFont(FONT_REG, 10)
        c2.drawString(2 * cm, 25 * cm, f"PDF konnte nicht eingebunden werden: {e}")
        c2.showPage()
        c2.save()
        sep2.seek(0)
        for p in PdfReader(sep2).pages:
            writer.add_page(p)
    return start_page


def build_gesamt_pdf(testakte_dir: Path) -> tuple[str, str]:
    """Gibt (status, info) zurueck. status in {ok, skip, error}."""
    name = testakte_dir.name
    out_dir = testakte_dir / "gesamt-pdf"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / f"{name}_gesamt.pdf"

    all_files = collect_files(testakte_dir)
    canonical_files = collect_files(testakte_dir, canonical_only=True)
    total_files = sum(len(v) for v in all_files.values())
    canonical_count = sum(len(v) for v in canonical_files.values())
    if canonical_count == 0:
        return "skip", "keine Quelldateien"

    einzel_count, einzel_errors = build_individual_pdfs(testakte_dir, canonical_files)
    if einzel_errors:
        preview = "; ".join(einzel_errors[:3])
        return "error", f"Einzel-PDFs konnten nicht erzeugt werden: {preview}"

    readme_h1, readme_summary = extract_readme_summary(testakte_dir / "README.md")
    cover = build_cover(name, readme_summary, readme_h1)

    defaults = TESTAKTE_DEFAULTS.get(name, {})
    tmp_handle = tempfile.NamedTemporaryFile(
        prefix=f".{name}-text-",
        suffix=".pdf",
        delete=False,
    )
    tmp_text = Path(tmp_handle.name)
    tmp_handle.close()
    output_handle = tempfile.NamedTemporaryFile(
        prefix=f".{name}-gesamt-",
        suffix=".pdf",
        dir=out_dir,
        delete=False,
    )
    tmp_output = Path(output_handle.name)
    output_handle.close()
    try:
        ok, pdf_attachments, bookmarks = build_text_pdf(
            testakte_dir,
            canonical_files,
            cover,
            tmp_text,
        )
        if not ok:
            return "error", "Text-PDF konnte nicht erzeugt werden"

        writer = PdfWriter()
        try:
            for page in PdfReader(str(tmp_text)).pages:
                writer.add_page(page)
        except Exception as e:
            return "error", f"Text-PDF nicht lesbar: {e}"

        for pdf in pdf_attachments:
            raw_text = raw_text_for_meta(pdf, "pdf")
            meta = document_meta(testakte_dir, pdf, raw_text)
            start_page = append_pdf_with_separator(
                writer,
                f"Originalanlage: {meta['betreff']}",
                pdf,
                name,
            )
            bookmarks.append((outline_label(pdf, meta), start_page))

        writer.add_outline_item("Aktenvorblatt", 0, bold=True)
        if bookmarks:
            parent = writer.add_outline_item(
                "Aktenstücke und Anlagen",
                bookmarks[0][1],
                bold=True,
            )
            for label, page_number in bookmarks:
                writer.add_outline_item(label, page_number, parent=parent)
            writer.page_mode = "/UseOutlines"

        writer.add_metadata(
            {
                "/Title": defaults.get("vorhaben", name),
                "/Author": defaults.get("auftraggeber", "Vergabeakte"),
                "/Subject": f"Vergabeakte | Az. {defaults.get('aktenzeichen', name)}",
                "/Keywords": f"Vergabeakte, {defaults.get('aktenzeichen', name)}",
            }
        )
        with open(tmp_output, "wb") as f:
            writer.write(f)
        tmp_output.replace(out_path)
    except Exception as e:
        return "error", f"Gesamt-PDF konnte nicht geschrieben werden: {e}"
    finally:
        tmp_text.unlink(missing_ok=True)
        tmp_output.unlink(missing_ok=True)
    size_kb = out_path.stat().st_size / 1024
    return "ok", (
        f"{out_path.relative_to(REPO_ROOT)} ({size_kb:.0f} KB, "
        f"{canonical_count} Aktenstücke/Anlagen, {einzel_count} Einzel-PDFs, "
        f"{total_files - canonical_count} Formatfassungen nicht doppelt eingebunden)"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Einzel- und Gesamt-PDFs der Testakten erzeugen."
    )
    parser.add_argument(
        "testakten",
        nargs="*",
        metavar="NAME",
        help="Testakten-Verzeichnis; ohne Angabe werden alle Akten gebaut.",
    )
    parser.add_argument(
        "--list",
        action="store_true",
        help="Verfügbare Testakten anzeigen und beenden.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    all_dirs = sorted(
        d for d in TESTAKTEN.iterdir()
        if d.is_dir() and d.name not in HELPER_DIRS
    )
    by_name = {d.name: d for d in all_dirs}
    if args.list:
        print("\n".join(by_name))
        return 0
    unknown = sorted(set(args.testakten) - set(by_name))
    if unknown:
        print(
            "FEHLER: unbekannte Testakte(n): " + ", ".join(unknown),
            file=sys.stderr,
        )
        return 2
    selected = [by_name[name] for name in args.testakten] if args.testakten else all_dirs
    print(f"Verarbeite {len(selected)} Testakten")
    print()
    counts = {"ok": 0, "skip": 0, "error": 0}
    for d in selected:
        status, info = build_gesamt_pdf(d)
        counts[status] += 1
        sigil = {"ok": "OK ", "skip": "SK ", "error": "ERR"}[status]
        print(f"  {sigil} {d.name}: {info}")
    print()
    print(f"Fertig: {counts['ok']} OK, {counts['skip']} skip, {counts['error']} Fehler")
    return 1 if counts["error"] else 0


if __name__ == "__main__":
    sys.exit(main())
