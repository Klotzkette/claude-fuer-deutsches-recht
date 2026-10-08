#!/usr/bin/env python3
"""Prueft die lebensechte Testakten-Schicht im Repo.

Die Release-ZIP-Validierung stellt sicher, dass die Exportpakete keine
Markdown-Dateien enthalten. Dieser Validator geht eine Ebene frueher an:
Jede Testakte muss den deterministisch erzeugten Aktenbeifang im Arbeitsbaum
enthalten, und die zugehoerigen Einzel-PDFs muessen bereits gebaut sein.
"""

from __future__ import annotations

import csv
import re
import sys
from datetime import datetime
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path
from zoneinfo import ZoneInfo

try:
    from PIL import Image
except ImportError:  # pragma: no cover - requirements-dev installiert Pillow
    Image = None  # type: ignore[assignment]
try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - requirements-dev installiert openpyxl
    load_workbook = None  # type: ignore[assignment]
try:
    from docx import Document
except ImportError:  # pragma: no cover - requirements-dev installiert python-docx
    Document = None  # type: ignore[assignment]
try:
    from pptx import Presentation
except ImportError:  # pragma: no cover - requirements-dev installiert python-pptx
    Presentation = None  # type: ignore[assignment]
from pypdf import PdfReader


REPO_ROOT = Path(__file__).resolve().parent.parent
TESTAKTEN = REPO_ROOT / "testakten"
SKIP_DIRS = {"formatvorlagen-paradebeispiele", "megaprompts"}

BEIFANG = {
    "notizen/00_bearbeitungsnotizen_eingangskorb.txt": "notizen__00_bearbeitungsnotizen_eingangskorb.pdf",
    "daten/00_portal_aktivitaetsprotokoll.csv": "daten__00_portal_aktivitaetsprotokoll.pdf",
    "eml/00_interne_weiterleitung_aktenstand.eml": "eml__00_interne_weiterleitung_aktenstand.pdf",
    "bilder/00_fristen_upload_skizze.png": "bilder__00_fristen_upload_skizze.pdf",
}

PORTAL_HEADER = [
    "Zeitstempel",
    "System",
    "Nutzer",
    "Vorgang",
    "Dokument",
    "Status",
    "SHA256",
    "Bemerkung",
]

OBSOLETE_ARTIFACT_STEMS = {
    "04-it-beratung-ministerium-musterland": {
        "06_olg_beschwerde_skizze",
    },
    "06-konkurrentenrechtsschutz-rechenzentrum-musterkreis": {
        "12_sofortige_beschwerde_olg",
    },
    "it-sig-2-vergabe-landeshauptstadt-schwerin-nachpruefung": {
        "17_beschluss_vk_bund",
        "18_sofortige_beschwerde_olg_duesseldorf",
        "20_de_facto_vergabe_klage",
    },
}

VISIBLE_TRANSLITERATIONS = re.compile(
    r"\b(?:angekuendigt|Vorlaeufig\w*|Supportkraeft\w*|Verguetungsmodell|"
    r"Vertragsaenderung|Aenderungen|Eignungspruefung|Urspruenglich|Federfuehrung|"
    r"Stueck|Unterstuetzung|Endgeraete|Servicequalitaet|Qualitaetskontrolle|"
    r"Klassenraeume|taeglich|Sanitaer|Maengel|Glasflaechen|jaehrlich|"
    r"Qualitaetssicherung|Ruegefrist|anhaengig|Ausloeser|ausgeloest|fruehestens|"
    r"ueberwachen|Rueckversetzung|oeffentliche|Unzulaessige|Nachaufklaerung|"
    r"Betriebsstaette|Geschaeftsgeheimnisse|Schwaerzung|Fristverlaengerung|"
    r"veroeffentlicht|OEPNV|UNZULAESSIGE|VERSCHAERFUNG|RUEGEFORDERUNG|GRUENDE|"
    r"haelftige)\b|Amt fuer|Finanzbehoerde|Grundschule Sued|KommunalCloud Sued|"
    r"Mit freundlichen Gruessen|Sabine Mueller"
)
BERLIN_TZ = ZoneInfo("Europe/Berlin")
MIN_LAST_PAGE_CHARS = 400
MIN_SEGMENT_LAST_PAGE_CHARS = 400
INTERNAL_SOURCE_SUFFIX = re.compile(
    r"\.(?:md|txt|csv|xml|x83|x84|d83|d84|g48|png|jpg|jpeg)$",
    re.IGNORECASE,
)


def fail(message: str) -> None:
    print(f"validate-testakten-echtformate failed: {message}", file=sys.stderr)
    raise SystemExit(1)


def assert_file(path: Path, *, min_bytes: int = 1) -> None:
    if not path.is_file():
        fail(f"{path}: missing")
    if path.stat().st_size < min_bytes:
        fail(f"{path}: too small")


def validate_note(path: Path) -> None:
    text = path.read_text(encoding="utf-8")
    required = ["Lose Bearbeitungsnotizen", "Frist", "Upload", "Portal"]
    missing = [needle for needle in required if needle not in text]
    if missing:
        fail(f"{path}: missing note markers {missing}")
    if any(marker in text.lower() for marker in ("demonstrationsakte", "plugin-test", "github-release")):
        fail(f"{path}: contains didactic/meta marker")


def source_document_dates(testakte_dir: Path) -> dict[str, datetime]:
    result = {}
    for path in sorted(testakte_dir.glob("[0-9]*.md")):
        text = path.read_text(encoding="utf-8")
        metadata = re.search(r"<!--\s*aktenmeta.*?^Datum:\s*(\d{2}\.\d{2}\.\d{4})\s*$.*?-->", text, re.S | re.M)
        match = metadata or re.search(r"\b(\d{2}\.\d{2}\.\d{4})\b", text)
        if match:
            value = datetime.strptime(match.group(1), "%d.%m.%Y")
            result[path.name] = value
            for suffix in (".docx", ".eml", ".pdf"):
                result[path.stem + suffix] = value
    return result


def validate_portal_csv(path: Path) -> None:
    with path.open(encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    if len(rows) < 9:
        fail(f"{path}: expected header plus at least 8 activity rows")
    if rows[0] != PORTAL_HEADER:
        fail(f"{path}: unexpected header {rows[0]!r}")
    source_dates = source_document_dates(path.parent.parent)
    previous_stamp: datetime | None = None
    for row_no, row in enumerate(rows[1:], start=2):
        if len(row) != len(PORTAL_HEADER):
            fail(f"{path}:{row_no}: expected {len(PORTAL_HEADER)} columns")
        if row[5] not in {"OK", "nachbearbeitet"}:
            fail(f"{path}:{row_no}: unexpected status {row[5]!r}")
        if not re.fullmatch(r"[0-9a-f]{16}", row[6]):
            fail(f"{path}:{row_no}: invalid short hash {row[6]!r}")
        try:
            stamp = datetime.strptime(row[0], "%Y-%m-%d %H:%M:%S")
        except ValueError:
            fail(f"{path}:{row_no}: invalid timestamp {row[0]!r}")
        if previous_stamp and stamp < previous_stamp:
            fail(f"{path}:{row_no}: activity log is not chronological")
        previous_stamp = stamp
        source_date = source_dates.get(row[4])
        if source_date and stamp.date() < source_date.date():
            fail(f"{path}:{row_no}: activity predates source document {row[4]}")


def validate_csv_shape(path: Path) -> None:
    with path.open(encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        fail(f"{path}: empty CSV")
    expected = len(rows[0])
    if expected == 0:
        fail(f"{path}: empty CSV header")
    for row_no, row in enumerate(rows[1:], start=2):
        if len(row) != expected:
            fail(f"{path}:{row_no}: {len(row)} columns, expected {expected}; quote embedded commas")


def validate_eml(path: Path, *, require_internal_markers: bool = False) -> None:
    raw = path.read_bytes()
    separator = b"\r\n\r\n" if b"\r\n\r\n" in raw else b"\n\n"
    raw_headers = raw.partition(separator)[0]
    if any(byte >= 128 for byte in raw_headers):
        fail(f"{path}: non-ASCII header is not RFC 2047 encoded")
    if any(len(line) > 998 for line in raw.replace(b"\r\n", b"\n").split(b"\n")):
        fail(f"{path}: line exceeds the RFC 5322 hard limit")
    msg = BytesParser(policy=policy.default).parsebytes(raw)
    defects = [type(defect).__name__ for part in msg.walk() for defect in part.defects]
    if defects:
        fail(f"{path}: malformed MIME message ({', '.join(defects)})")
    for part in msg.walk():
        if str(part.get("Content-Transfer-Encoding", "")).lower() != "quoted-printable":
            continue
        encoded = part.get_payload()
        if isinstance(encoded, str) and any(len(line) > 76 for line in encoded.splitlines()):
            fail(f"{path}: quoted-printable line exceeds 76 characters")
    for header in ("From", "To", "Subject", "Date", "Message-ID"):
        values = msg.get_all(header, [])
        if not values:
            fail(f"{path}: missing {header}")
        if len(values) != 1:
            fail(f"{path}: duplicate {header}")
    try:
        sent_at = parsedate_to_datetime(str(msg["Date"]))
    except (TypeError, ValueError, OverflowError):
        fail(f"{path}: invalid Date header")
    expected_offset = sent_at.replace(tzinfo=BERLIN_TZ).utcoffset()
    if sent_at.utcoffset() != expected_offset:
        fail(f"{path}: Date header has the wrong Europe/Berlin UTC offset")
    if not re.fullmatch(r"<[^<>\s@]+@[^<>\s@]+>", str(msg["Message-ID"])):
        fail(f"{path}: invalid Message-ID")
    body_part = msg.get_body(preferencelist=("plain",))
    if body_part is None:
        fail(f"{path}: missing plain-text body")
    body = body_part.get_content()
    if not isinstance(body, str) or len(body.strip()) < 40:
        fail(f"{path}: plain-text body is empty or implausibly short")
    match = VISIBLE_TRANSLITERATIONS.search(f"{msg.get('Subject', '')}\n{body}")
    if match:
        fail(f"{path}: visible ASCII transliteration {match.group(0)!r}")
    combined = f"{msg.get('Subject', '')}\n{body}"
    if require_internal_markers and ("Akte:" not in body or "Unterlagen" not in combined):
        fail(f"{path}: internal forwarding body is incomplete")


def validate_visible_text(path: Path) -> None:
    text = path.read_text(encoding="utf-8")
    match = VISIBLE_TRANSLITERATIONS.search(text)
    if match:
        fail(f"{path}: visible ASCII transliteration {match.group(0)!r}")
    if re.search(r"\bDokument-Ende\b", text, re.IGNORECASE):
        fail(f"{path}: contains an internal document-production marker")


def validate_no_internal_source_references(testakte_dir: Path) -> None:
    patterns = ("*.csv", "*.txt", "*.eml", "*.xml", "*.g48", "*.x83", "*.x84", "*.d83", "*.d84")
    for pattern in patterns:
        for path in sorted(testakte_dir.rglob(pattern)):
            text = path.read_text(encoding="utf-8", errors="replace")
            match = re.search(r"(?<![\w.-])[\w.-]+\.md\b", text, re.IGNORECASE)
            if match:
                fail(f"{path}: internal Markdown source reference {match.group(0)!r}")


def validate_ancillary_timeline(testakte_dir: Path) -> None:
    dates = source_document_dates(testakte_dir)
    if not dates:
        return
    latest = max(dates.values()).date()
    note = (testakte_dir / "notizen" / "00_bearbeitungsnotizen_eingangskorb.txt").read_text(encoding="utf-8")
    expected = latest.strftime("%d.%m.%Y")
    if f"Stand: {expected}" not in note:
        fail(f"{testakte_dir}: loose-note status does not match latest source date {expected}")
    internal = testakte_dir / "eml" / "00_interne_weiterleitung_aktenstand.eml"
    msg = BytesParser(policy=policy.default).parsebytes(internal.read_bytes())
    sent = parsedate_to_datetime(str(msg["Date"])).date()
    if sent <= latest:
        fail(f"{internal}: forwarding date must follow latest source date {expected}")


def validate_png(path: Path) -> None:
    assert_file(path, min_bytes=10_000)
    if Image is None:
        return
    with Image.open(path) as img:
        if img.size[0] < 1000 or img.size[1] < 700:
            fail(f"{path}: image too small {img.size}")
        if img.mode not in {"RGB", "RGBA"}:
            fail(f"{path}: unexpected image mode {img.mode}")


def validate_markdown_pdf_pagination(testakte_dir: Path) -> None:
    for source in sorted(testakte_dir.glob("[0-9][0-9]_*.md")):
        pdf = testakte_dir / "einzel-pdf" / f"{source.stem}.pdf"
        assert_file(pdf, min_bytes=20_000)
        try:
            reader = PdfReader(str(pdf))
            if len(reader.pages) <= 1:
                continue
            text = reader.pages[-1].extract_text() or ""
        except Exception as exc:
            fail(f"{pdf}: cannot inspect pagination ({exc})")
        chars = len(re.sub(r"\s+", "", text))
        if chars < MIN_LAST_PAGE_CHARS:
            fail(
                f"{pdf}: sparse final page with only {chars} non-whitespace characters"
            )


def outline_entry_count(items: list) -> int:
    count = 0
    for item in items:
        if isinstance(item, list):
            count += outline_entry_count(item)
        else:
            count += 1
    return count


def validate_pdf_properties(testakte_dir: Path) -> None:
    slug = testakte_dir.name
    gesamt = testakte_dir / "gesamt-pdf" / f"{slug}_gesamt.pdf"
    try:
        reader = PdfReader(str(gesamt))
    except Exception as exc:
        fail(f"{gesamt}: cannot read PDF ({exc})")
    metadata = reader.metadata
    title = (metadata.title or "").strip()
    subject = (metadata.subject or "").strip()
    author = (metadata.author or "").strip()
    if not title or title in {slug, f"Arbeitsakte {slug}"}:
        fail(f"{gesamt}: technical or missing PDF title")
    if not subject or "Az." not in subject:
        fail(f"{gesamt}: subject must identify the Vergabeakte and Aktenzeichen")
    if not author or author == "Vergabeakte":
        fail(f"{gesamt}: author must identify the responsible organisation")
    individual = sorted((testakte_dir / "einzel-pdf").glob("*.pdf"))
    expected_outline_entries = len(individual) + 2
    actual_outline_entries = outline_entry_count(reader.outline)
    if actual_outline_entries != expected_outline_entries:
        fail(
            f"{gesamt}: {actual_outline_entries} outline entries, "
            f"expected {expected_outline_entries}"
        )
    if str(reader.root_object.get("/PageMode", "")) != "/UseOutlines":
        fail(f"{gesamt}: PDF viewer is not instructed to open the outline")
    if len(reader.outline) < 3 or not isinstance(reader.outline[2], list):
        fail(f"{gesamt}: Aktenstückgliederung im PDF-Lesezeichenbaum fehlt")
    destinations = reader.outline[2]
    starts = [reader.get_destination_page_number(item) for item in destinations]
    if any(page < 0 for page in starts) or starts != sorted(starts):
        fail(f"{gesamt}: Aktenstück-Lesezeichen sind nicht chronologisch")
    for index, start in enumerate(starts):
        end = starts[index + 1] - 1 if index + 1 < len(starts) else len(reader.pages) - 1
        if end <= start:
            continue
        start_text = reader.pages[start].extract_text() or ""
        if "Originalanlage" in start_text:
            continue
        end_text = reader.pages[end].extract_text() or ""
        chars = len(re.sub(r"\s+", "", end_text))
        if chars < MIN_SEGMENT_LAST_PAGE_CHARS:
            title = getattr(destinations[index], "title", f"Aktenstück {index + 1}")
            fail(
                f"{gesamt}: {title!r} endet auf einer dünn belegten Seite "
                f"mit nur {chars} Zeichen"
            )

    for pdf in individual:
        try:
            item_reader = PdfReader(str(pdf))
        except Exception as exc:
            fail(f"{pdf}: cannot read PDF ({exc})")
        item_metadata = item_reader.metadata
        item_title = (item_metadata.title or "").strip()
        item_subject = (item_metadata.subject or "").strip()
        item_author = (item_metadata.author or "").strip()
        if (
            not item_title
            or item_title.startswith("Einzel-PDF ")
            or INTERNAL_SOURCE_SUFFIX.search(item_title)
        ):
            fail(f"{pdf}: technical or missing PDF title {item_title!r}")
        if not item_subject or item_subject in {"(unspecified)", "Einzelakte"}:
            fail(f"{pdf}: technical or missing PDF subject")
        if "Az." not in item_subject:
            fail(f"{pdf}: PDF subject does not contain the Aktenzeichen")
        if not item_author or item_author in {"Vergabeakte", "Kanzleiakte"}:
            fail(f"{pdf}: technical or missing PDF author")
        first_page = item_reader.pages[0].extract_text() or ""
        if f"Arbeitsakte: {slug}" in first_page:
            fail(f"{pdf}: technical directory slug is visible in the footer")
        if re.search(r"Quelle:\s*.*\.(?:md|txt|csv|xml)", first_page, re.I | re.S):
            fail(f"{pdf}: internal source filename is visible in the document header")


def validate_docx_properties(path: Path) -> None:
    if Document is None:
        return
    document = Document(str(path))
    props = document.core_properties
    if not props.title or INTERNAL_SOURCE_SUFFIX.search(props.title):
        fail(f"{path}: missing or technical DOCX title")
    if not props.subject or "Az." not in props.subject:
        fail(f"{path}: DOCX subject does not contain the Aktenzeichen")
    if not props.author or props.author == "python-docx":
        fail(f"{path}: DOCX author is missing or technical")
    if not props.last_modified_by or props.last_modified_by == "python-docx":
        fail(f"{path}: DOCX last-modified identity is missing or technical")


def validate_pptx_properties(path: Path) -> None:
    if Presentation is None:
        return
    presentation = Presentation(str(path))
    props = presentation.core_properties
    if not props.title or INTERNAL_SOURCE_SUFFIX.search(props.title):
        fail(f"{path}: missing or technical PPTX title")
    if not props.subject or "Az." not in props.subject:
        fail(f"{path}: PPTX subject does not contain the Aktenzeichen")
    if not props.author:
        fail(f"{path}: PPTX author is missing")


def validate_scan_pdf_properties(path: Path) -> None:
    reader = PdfReader(str(path))
    metadata = reader.metadata
    title = (metadata.title or "").strip()
    subject = (metadata.subject or "").strip()
    author = (metadata.author or "").strip()
    if not title or title == path.stem or "_" in title:
        fail(f"{path}: scan PDF title is missing or technical")
    if not subject or "Az." not in subject:
        fail(f"{path}: scan PDF subject does not contain the Aktenzeichen")
    if not author:
        fail(f"{path}: scan PDF author is missing")


def validate_generated_xlsx(path: Path) -> None:
    if load_workbook is None:
        return
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        properties = workbook.properties
        if not properties.title or INTERNAL_SOURCE_SUFFIX.search(properties.title):
            fail(f"{path}: missing or technical XLSX title")
        if not properties.subject or "Az." not in properties.subject:
            fail(f"{path}: XLSX subject does not contain the Aktenzeichen")
        if not properties.creator or properties.creator == "openpyxl":
            fail(f"{path}: XLSX creator is missing or technical")
        if not properties.lastModifiedBy or properties.lastModifiedBy == "openpyxl":
            fail(f"{path}: XLSX last-modified identity is missing or technical")
        if "Quelle" not in workbook.sheetnames:
            return
        if workbook.sheetnames[0] != "Quelle":
            fail(f"{path}: generated workbook must start with the Quelle sheet")
        table_sheets = [
            sheet for sheet in workbook.worksheets if sheet.title.startswith("Tabelle ")
        ]
        if not table_sheets:
            fail(f"{path}: generated workbook has no table sheet")
        for number, sheet in enumerate(table_sheets, start=1):
            if sheet.title != f"Tabelle {number}":
                fail(f"{path}: non-contiguous table sheet numbering")
            if sheet.print_title_rows != "$1:$1":
                fail(f"{path}/{sheet.title}: first row is not repeated when printing")
            if sheet.page_setup.fitToWidth != 1:
                fail(f"{path}/{sheet.title}: print layout does not fit to one page width")
            if not sheet.sheet_properties.pageSetUpPr.fitToPage:
                fail(f"{path}/{sheet.title}: fit-to-page is disabled")
            margins = sheet.page_margins
            if margins.top < 0.7 or margins.header > 0.3:
                fail(f"{path}/{sheet.title}: header can collide with table content")
            if margins.bottom < 0.7 or margins.footer < 0.3:
                fail(f"{path}/{sheet.title}: footer can be clipped")
            header = sheet.oddHeader.center.text or ""
            if not header.endswith(f"| Tabelle {number}"):
                fail(f"{path}/{sheet.title}: concise print header is missing")
            if sheet.oddFooter.right.text != "Seite &P von &N":
                fail(f"{path}/{sheet.title}: print pagination is missing")
    finally:
        workbook.close()


def validate_testakte(testakte_dir: Path) -> None:
    gesamt_pdf = testakte_dir / "gesamt-pdf" / f"{testakte_dir.name}_gesamt.pdf"
    assert_file(gesamt_pdf, min_bytes=20_000)

    for rel, pdf_name in BEIFANG.items():
        path = testakte_dir / rel
        assert_file(path)
        einzel_pdf = testakte_dir / "einzel-pdf" / pdf_name
        assert_file(einzel_pdf, min_bytes=20_000)

    validate_note(testakte_dir / "notizen" / "00_bearbeitungsnotizen_eingangskorb.txt")
    validate_portal_csv(testakte_dir / "daten" / "00_portal_aktivitaetsprotokoll.csv")
    internal_eml = testakte_dir / "eml" / "00_interne_weiterleitung_aktenstand.eml"
    for eml in sorted((testakte_dir / "eml").glob("*.eml")):
        validate_eml(eml, require_internal_markers=eml == internal_eml)
    for suffix in ("*.md", "*.csv", "*.xml", "*.txt"):
        for text_file in sorted(testakte_dir.rglob(suffix)):
            validate_visible_text(text_file)
            if text_file.suffix == ".csv":
                validate_csv_shape(text_file)
    validate_png(testakte_dir / "bilder" / "00_fristen_upload_skizze.png")
    validate_no_internal_source_references(testakte_dir)
    validate_ancillary_timeline(testakte_dir)
    validate_markdown_pdf_pagination(testakte_dir)
    validate_pdf_properties(testakte_dir)
    for docx_path in sorted((testakte_dir / "docx").glob("*.docx")):
        validate_docx_properties(docx_path)
    for xlsx in sorted((testakte_dir / "xlsx").glob("*.xlsx")):
        validate_generated_xlsx(xlsx)
    for pptx_path in sorted((testakte_dir / "pptx").glob("*.pptx")):
        validate_pptx_properties(pptx_path)
    for scan_pdf in sorted((testakte_dir / "scan-pdf").glob("*.pdf")):
        validate_scan_pdf_properties(scan_pdf)

    for stem in OBSOLETE_ARTIFACT_STEMS.get(testakte_dir.name, set()):
        for subdir, suffix in (
            ("docx", ".docx"),
            ("eml", ".eml"),
            ("xlsx", ".xlsx"),
            ("scan-pdf", ".pdf"),
            ("pptx", ".pptx"),
            ("einzel-pdf", ".pdf"),
        ):
            obsolete = testakte_dir / subdir / f"{stem}{suffix}"
            if obsolete.exists():
                fail(f"{obsolete}: obsolete artifact from renamed source")
        einzel_dir = testakte_dir / "einzel-pdf"
        for candidate in einzel_dir.glob("*.pdf"):
            if candidate.stem.endswith(f"__{stem}"):
                fail(f"{candidate}: obsolete derived PDF from renamed source")


def main() -> None:
    dirs = sorted(d for d in TESTAKTEN.iterdir() if d.is_dir() and d.name not in SKIP_DIRS)
    if not dirs:
        fail("no testakten directories found")
    for testakte_dir in dirs:
        validate_testakte(testakte_dir)
    print(f"validate-testakten-echtformate OK ({len(dirs)} Testakten)")


if __name__ == "__main__":
    main()
