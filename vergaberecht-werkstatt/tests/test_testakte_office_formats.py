from __future__ import annotations

import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from pptx import Presentation


ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "build-testakten-echtformate.py"
sys.path.insert(0, str(SCRIPT.parent))
SPEC = importlib.util.spec_from_file_location("build_testakten_echtformate", SCRIPT)
assert SPEC is not None and SPEC.loader is not None
formats = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(formats)

META = {
    "absender": "Stadt Musterstadt - Zentrale Vergabestelle",
    "adresse": "Rathausplatz 1, 99999 Musterstadt",
    "empfaenger": "Vergabeakte",
    "datum": "26.07.2026",
    "aktenzeichen": "VS-2026-TEST",
    "dokumenttyp": "Prüfvermerk",
    "betreff": "Prüfvermerk zur Ausgabequalität",
}


class TestakteOfficeFormatTests(unittest.TestCase):
    def test_only_explicit_handcrafted_counterpart_suppresses_generation(self) -> None:
        case_name = "it-sig-2-vergabe-landeshauptstadt-schwerin-nachpruefung"
        with tempfile.TemporaryDirectory() as directory:
            case_dir = Path(directory) / case_name
            (case_dir / "xlsx").mkdir(parents=True)
            (case_dir / "xlsx" / "09_nachpruefungsantrag_vk_mv.xlsx").touch()
            source = case_dir / "09_nachpruefungsantrag_vk_mv.md"
            self.assertFalse(formats.hat_handgefertigtes_gegenstueck(case_dir, source))
            counterpart = case_dir / "docx" / "nachpruefungsantrag_vk.docx"
            counterpart.parent.mkdir()
            counterpart.touch()
            self.assertTrue(formats.hat_handgefertigtes_gegenstueck(case_dir, source))

    def test_docx_tail_is_kept_together(self) -> None:
        blocks = [
            ("heading", (1, META["betreff"])),
            ("heading", (2, "Prüfung")),
            ("para", "A" * 500),
            ("heading", (2, "Entscheidung")),
            ("para", "B" * 500),
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "pruefvermerk.docx"
            formats.write_docx(path, META, blocks)
            document = Document(path)
        nonempty = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        self.assertEqual(document.styles["Normal"].font.size.pt, 10.5)
        self.assertTrue(nonempty[-1].paragraph_format.keep_together)
        self.assertTrue(nonempty[-2].paragraph_format.keep_with_next)

    def test_pptx_tables_use_context_titles_and_split_safely(self) -> None:
        rows = [["Kriterium", "Bewertung"]]
        rows.extend([[f"Kriterium {index}", f"Bewertung {index}"] for index in range(14)])
        blocks = [
            ("heading", (1, META["betreff"])),
            ("heading", (2, "Wertungsmatrix")),
            ("table", rows),
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "pruefvermerk.pptx"
            formats.write_pptx(path, META, blocks)
            presentation = Presentation(path)
        titles = [
            slide.shapes.title.text
            for slide in presentation.slides
            if slide.shapes.title is not None
        ]
        self.assertIn("Wertungsmatrix (1/2)", titles)
        self.assertIn("Wertungsmatrix (2/2)", titles)
        self.assertNotIn("Tabelle", titles)
        table_rows = [
            len(shape.table.rows)
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_table
        ]
        self.assertEqual(table_rows, [13, 3])


if __name__ == "__main__":
    unittest.main()
