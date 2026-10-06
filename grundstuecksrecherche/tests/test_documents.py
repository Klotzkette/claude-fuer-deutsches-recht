"""Isolierte Modultests mit technischen Eingaben, keine Testakten oder Liveabrufe."""

from copy import deepcopy
from html.parser import HTMLParser
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import socket
import sys
import unittest
from unittest.mock import patch
from zipfile import ZipFile, is_zipfile

from docx import Document
from docx.oxml.ns import qn


MODULE_PATH = Path(__file__).resolve().parents[1] / "app" / "documents.py"
SPEC = importlib.util.spec_from_file_location("parcel_documents_under_test", MODULE_PATH)
documents = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = documents
SPEC.loader.exec_module(documents)


def parcel(index=0):
    # Eindeutig als technische Prüfdaten erkennbar; keine behauptete reale Fläche.
    return {
        "stable_id": f"test-selection-{index}", "provider_id": "unit-test",
        "municipality": f"Prüfgemeinde-{index}", "municipality_code": "05515000",
        "district_name": f"Prüfgemarkung-{index}", "district_code": "test-district",
        "flur": "007", "numerator": f"{index + 1:05}", "denominator": "0002",
        "official_parcel_reference": f"TEST-KEIN-AMTLICHES-KENNZEICHEN-{index}",
        "geometry": {"type": "Polygon", "coordinates": [[[7, 51], [7.1, 51], [7.1, 51.1], [7, 51]]]},
        "area_value": 123.5, "area_unit": "m²", "location_text": f"Technischer Lagehinweis {index}",
        "source_url": "https://www.stadt-muenster.de/katasteramt/startseite",
        "source_date": "2026-10-01", "retrieved_at": "2026-10-06T09:00:00Z",
        "identification_status": "manuell_ergaenzt",
    }


def case_data(count=2, state="Nordrhein-Westfalen"):
    return {
        "case_id": "UNIT-TEST", "profile": {
            "name": "Prüfgemeinde", "state": state, "municipality_code": "05515000",
            "center": [51, 7], "bounds": [6, 50, 8, 52],
            "authorities": [], "providers": [], "evidence": [],
        },
        "selected_parcels": [parcel(index) for index in range(count)],
        "purpose": "Die technische Ausgabe soll geprüft werden.",
        "specific_interest": "Eine rechtliche Berechtigung wird nicht behauptet.",
        "requested_information_scope": "Es wird nur die technische Ausgabe geprüft.",
        "sender_organisation": "Technischer Absender", "sender_address": "Keine echte Anschrift",
        "contact_person": "Prüfkontakt", "legal_department_contact": "Prüfstelle Recht",
        "recipients": {"kataster": "Prüfempfänger Kataster", "grundbuch": "Prüfempfänger Grundbuch", "notar": "Prüfempfänger Notar"},
        "attachments": "Keine beigefügten Dateien", "created_at": "2026-10-01", "updated_at": "2026-10-06",
    }


class ParsedHTML(HTMLParser):
    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.tags = []
        self.attributes = []
        self.texts = []
        self.paragraphs = []
        self._paragraph = None
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        self.attributes.extend(attrs)
        if tag == "p":
            self._paragraph = []

    def handle_endtag(self, tag):
        if tag == "p" and self._paragraph is not None:
            self.paragraphs.append("".join(self._paragraph))
            self._paragraph = None

    def handle_data(self, text):
        self.texts.append(text)
        if self._paragraph is not None:
            self._paragraph.append(text)

    @property
    def text(self):
        return "\n".join(self.texts)


class ValidationTests(unittest.TestCase):
    def test_empty_case_and_nullable_profile_are_valid(self):
        for value in ({}, {"profile": None}, {"selected_parcels": []}):
            self.assertEqual(documents.validate_case(value), value)

    def test_copy_preserves_all_values_without_mutation(self):
        case = case_data()
        before = deepcopy(case)
        result = documents.validate_case(case)
        result["selected_parcels"][0]["geometry"]["coordinates"][0][0][0] = 8
        self.assertEqual(case, before)
        self.assertEqual(result["selected_parcels"][0]["numerator"], "00001")

    def test_malformed_roots_and_nested_types(self):
        bad_values = [None, [], "{}", b"{}", 1, True, {"profile": []},
                      {"profile": {"name": 1}}, {"selected_parcels": {}},
                      {"selected_parcels": [None]}, {"recipients": {"kataster": []}},
                      {"attachments": []}, {"case_id": None},
                      {"selected_parcels": [{"numerator": 12}]},
                      {"profile": {"evidence": ["https://example.org"]}},
                      {"profile": {"authorities": [{"evidence_urls": "https://example.org"}]}},
                      {"profile": {"providers": [{"crs": {"value": "EPSG:4326"}}]}}]
        for value in bad_values:
            with self.subTest(value=value), self.assertRaises(ValueError):
                documents.validate_case(value)

    def test_unknown_fields_rejected_at_every_level(self):
        invalid = [{"__proto__": {}}, {"send": True}, {"output_path": "/tmp/leak"},
                   {"profile": {"admin": True}}, {"selected_parcels": [{"owner": "Unknown"}]},
                   {"recipients": {"email": "unknown@example.org"}},
                   {"profile": {"providers": [{"file": "/etc/passwd"}]}},
                   {"profile": {"authorities": [{"verified": True}]}},
                   {"profile": {"evidence": [{"script": "run"}]}}]
        for value in invalid:
            with self.subTest(value=value), self.assertRaises(ValueError):
                documents.validate_case(value)

    def test_200_parcel_limit_and_duplicate_ids(self):
        self.assertEqual(len(documents.validate_case(case_data(200))["selected_parcels"]), 200)
        with self.assertRaises(ValueError):
            documents.validate_case(case_data(201))
        with self.assertRaises(ValueError):
            documents.validate_case({"selected_parcels": [parcel(), parcel()]})

    def test_2mb_limit_counts_utf8_bytes(self):
        case = case_data(200)
        for item in case["selected_parcels"]:
            item["location_text"] = "ä" * 6000
        with self.assertRaisesRegex(ValueError, "2 MiB"):
            documents.validate_case(case)

    def test_invalid_text_characters_and_deep_or_cyclic_input(self):
        for text in ("\x00", "\ud800", "\uffff", "a" * 20001):
            with self.subTest(text=repr(text[:10])), self.assertRaises(ValueError):
                documents.validate_case({"purpose": text})
        case = {}
        case["profile"] = case
        with self.assertRaises(ValueError):
            documents.validate_case(case)

    def test_non_finite_and_boolean_area_rejected(self):
        for number in (float("nan"), float("inf"), -float("inf"), True, -1, "123", "", 10 ** 400):
            with self.subTest(number=str(number)), self.assertRaises(ValueError):
                documents.validate_case({"selected_parcels": [{"area_value": number}]})

    def test_url_scheme_credentials_and_local_hosts_rejected(self):
        invalid = ["javascript:alert(1)", "data:text/html,bad", "file:///etc/passwd", "//example.org/a",
                   "ftp://example.org", "https://user:pass@example.org/a", "http://localhost/a",
                   "http://127.0.0.1", "http://[::1]", "http://10.0.0.1", "http://host.internal/",
                   "https://example.org\n/a", "https://example.org\\@localhost", "http://[", "http://example.org:99999"]
        for url in invalid:
            with self.subTest(url=url), self.assertRaises(ValueError):
                documents.validate_case({"selected_parcels": [{"source_url": url}]})
        for url in ("", "https://www.gesetze-im-internet.de/gbo/__12.html", "http://example.org/?a=1&b=2"):
            self.assertEqual(documents.validate_case({"selected_parcels": [{"source_url": url}]}),
                             {"selected_parcels": [{"source_url": url}]})

    def test_all_profile_source_urls_are_checked(self):
        for section, key in (("providers", "service_url"), ("providers", "tested_url"),
                             ("authorities", "official_form_url"), ("authorities", "official_website"),
                             ("evidence", "url"), ("evidence", "source_url")):
            with self.subTest(section=section, key=key), self.assertRaises(ValueError):
                documents.validate_case({"profile": {section: [{key: "file:///etc/passwd"}]}})

    def test_profile_claims_are_downgraded_even_with_future_dates(self):
        case = case_data()
        for key in ("providers", "authorities", "evidence"):
            case["profile"][key] = [{"status": "verifiziert", "verification_state": "verified", "verified_at": "2099-01-01"}]
        checked = documents.validate_case(case)
        for key in ("providers", "authorities", "evidence"):
            self.assertEqual(checked["profile"][key][0]["status"], "gefunden_ungeprueft")
            self.assertEqual(checked["profile"][key][0]["verification_state"], "gefunden_ungeprueft")
            self.assertEqual(case["profile"][key][0]["status"], "verifiziert")
        self.assertEqual(documents.validate_case(checked), checked)

    def test_manual_and_unavailable_status_remain_distinct(self):
        for status in ("manuell_ergaenzt", "nicht_verfuegbar", "eingeschraenkt"):
            case = {"profile": {"evidence": [{"status": status}]}}
            self.assertEqual(documents.validate_case(case), case)

    def test_malformed_geometries(self):
        invalid = [[], {}, {"type": "Point", "coordinates": [True, 51]},
                   {"type": "Point", "coordinates": [7, float("inf")]},
                   {"type": "Point", "coordinates": [181, 51]},
                   {"type": "Point", "coordinates": [7, 91]},
                   {"type": "Point", "coordinates": [7]},
                   {"type": "LineString", "coordinates": [[7, 51]]},
                   {"type": "Polygon", "coordinates": []},
                   {"type": "Polygon", "coordinates": [[[7, 51], [8, 51], [8, 52], [7, 52]]]},
                   {"type": "MultiPolygon", "coordinates": [[[]]]},
                   {"type": "GeometryCollection", "geometries": []},
                   {"type": "Point", "coordinates": [7, 51], "crs": "EPSG:4326"}]
        for geometry in invalid:
            with self.subTest(geometry=geometry), self.assertRaises(ValueError):
                documents.validate_case({"selected_parcels": [{"geometry": geometry}]})

    def test_multipolygons_holes_and_manual_points_preserved(self):
        outer = [[7, 51], [8, 51], [8, 52], [7, 51]]
        inner = [[7.1, 51.1], [7.2, 51.1], [7.2, 51.2], [7.1, 51.1]]
        for geometry in (None, {"type": "Point", "coordinates": [7, 51, 100]},
                         {"type": "MultiPolygon", "coordinates": [[outer, inner], [outer]]}):
            case = {"selected_parcels": [{"geometry": geometry}]}
            self.assertEqual(documents.validate_case(case), case)

    def test_geometry_byte_position_and_aggregate_budgets(self):
        for coordinates, message in (([[7, 51]] * 10001, "10000"),
                                     ([[7.123456789012345, 51.12345678901234]] * 9000, "256 KiB")):
            with self.subTest(message=message), self.assertRaisesRegex(ValueError, message):
                documents.validate_case({"selected_parcels": [{"geometry": {"type": "MultiPoint", "coordinates": coordinates}}]})
        with self.assertRaises(ValueError):
            documents.validate_case({"selected_parcels": [{"geometry": {"type": "MultiPoint", "coordinates": [[7, 51]] * 9000}} for _ in range(5)]})

    def test_invalid_centers_and_bounds(self):
        for profile in ({"center": [100, 7]}, {"center": {"latitude": 51}},
                        {"bounds": [8, 52, 6, 50]}, {"bounds": [6, 50, 8]}):
            with self.subTest(profile=profile), self.assertRaises(ValueError):
                documents.validate_case({"profile": profile})

    def test_real_discovery_contract_with_offline_responses(self):
        path = MODULE_PATH.with_name("discovery.py")
        if not path.exists():
            self.skipTest("Discovery-Modul wird separat implementiert.")
        spec = importlib.util.spec_from_file_location("discovery_contract_under_test", path)
        discovery = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(discovery)

        def search_response(url, **kwargs):
            feature = {"type": "Feature", "properties": {"gen": "Prüfgemeinde", "ags": "05515000"},
                       "bbox": [7.5, 51.8, 7.8, 52.1]}
            if "vg250_krs" in url:
                feature = {"type": "Feature", "properties": {"gen": "Prüfkreis", "ags": "05515"}}
            return json.dumps({"type": "FeatureCollection", "features": [feature]}).encode()

        def unavailable(url, **kwargs):
            raise TimeoutError("Bewusst simulierte, nicht live geprüfte Nichterreichbarkeit.")

        place = discovery.search_places("Prüfgemeinde", search_response)[0]
        found = discovery.discover(place, unavailable)
        profile = dict(place, **found, profile_id="contract-test")
        for provider in profile["providers"]:
            provider.update(source_url=provider["service_url"], verified_at="2026-10-06", test_result="Prüfeingabe", format="image/png")
        case = case_data()
        case["profile"] = profile
        checked = documents.validate_case(case)
        self.assertEqual(checked["profile"]["center"], place["center"])
        self.assertEqual(checked["profile"]["verification_state"], "gefunden_ungeprueft")
        self.assertEqual(len(documents.build_documents(case)), 4)
        self.assertTrue(is_zipfile(BytesIO(documents.export_document(case, "zip")[0])))

    def test_provider_metadata_contract_and_nested_rejections(self):
        profile = {
            "id": "ags:05515000", "profile_id": "test", "name": "Prüfgemeinde", "state": "Nordrhein-Westfalen",
            "state_code": "05", "district": "", "municipality_code": "05515000", "center": [51, 7],
            "bounds": [6, 50, 8, 52], "source_url": "https://example.org/places", "source_urls": [],
            "verification_state": "amtlich_verifiziert", "warnings": [],
            "providers": [{"provider_id": "test", "title": "Prüfdienst", "role": "parcels", "protocol": "WFS",
                           "type_name": "ave:Flurstueck", "layers": "", "crs": "EPSG:4326", "format": "application/json",
                           "license": "Metadatenangabe", "attribution": "Prüfquelle", "source_url": "https://example.org/wfs",
                           "verified_at": None, "test_result": "Prüfeingabe", "requires_host_approval": True,
                           "coverage": {"bounds": [[6, 50, 8, 52]], "verification_state": "metadatenangabe"}}],
            "authorities": [{"type": "grundbuch", "verified_at": None}],
            "evidence": [{"kind": "catalog_query", "source_url": "https://example.org/csw", "matched": None, "returned": "12"}],
        }
        self.assertEqual(len(documents.build_documents({"profile": profile})), 4)
        for path, value in (("license", {}), ("test_result", {}), ("requires_host_approval", "true"),
                            ("coverage", {"arbitrary": True}), ("layers", [1])):
            broken = deepcopy(profile)
            broken["providers"][0][path] = value
            with self.subTest(path=path), self.assertRaises(ValueError):
                documents.validate_case({"profile": broken})


class DocumentTests(unittest.TestCase):
    def test_undivided_parcel_does_not_request_an_unnecessary_denominator(self):
        case = case_data(1)
        case['selected_parcels'][0]['denominator'] = ''
        for document in documents.build_documents(case):
            self.assertIn('Flurstücksnummer mit 00001', document['html'])
            self.assertNotIn('Nenner mit [ergänzen]', document['html'])

    def test_exact_four_documents_and_complete_html(self):
        results = documents.build_documents({})
        self.assertEqual([item["id"] for item in results], ["uebergabe", "kataster", "grundbuch", "notar"])
        for item in results:
            self.assertEqual(set(item), {"id", "title", "html"})
            self.assertTrue(item["html"].startswith("<!DOCTYPE html>"))
            self.assertTrue(item["html"].endswith("</body></html>"))
            self.assertIn(documents.DRAFT_MARKER, item["html"])
            self.assertIn("[ergänzen]", item["html"])
            self.assertIn('font-size:11pt', item["html"])
            self.assertIn('Times New Roman', item["html"])
            self.assertNotIn("\u00a7", item["html"])

    def test_all_200_parcels_and_every_identifying_field_in_every_document(self):
        case = case_data(200)
        case["selected_parcels"][0].update({"grundbuchblatt": "TEST-BLATT-001", "grundbuchblaetter": ["TEST-BLATT-002"],
                                              "grundbuchbezirk": "Prüfbezirk", "grundbuchblatt_source": "Manueller Nachweis", "grundbuchblatt_date": "2026-09-30"})
        for item in documents.build_documents(case):
            text = ParsedHTML(item["html"]).text
            for selected in case["selected_parcels"]:
                for key, value in selected.items():
                    if key in ("geometry", "grundbuchblaetter"):
                        continue
                    self.assertIn(str(value), text, (item["id"], key))
            self.assertIn("TEST-BLATT-002", text)
            self.assertIn("manuell ergänzt", text)

    def test_missing_values_no_invented_numbers_or_owners(self):
        case = {"selected_parcels": [{"geometry": {"type": "Point", "coordinates": [7, 51]}, "identification_status": "lagehinweis"}]}
        for item in documents.build_documents(case):
            text = ParsedHTML(item["html"]).text
            self.assertIn("die Flurstücksnummer mit [ergänzen]", text)
            self.assertIn("Flurstückskennzeichen lautet nach den Vorgangsdaten [ergänzen]", text)
            self.assertIn("Lagehinweis", text)
            self.assertIn("kein Eigentum", text)
            self.assertNotIn("00000", text)

    def test_zero_area_is_not_missing(self):
        case = {"selected_parcels": [{"area_value": 0, "area_unit": "m²"}]}
        self.assertIn("Fläche beträgt 0 m²", documents.build_documents(case)[0]["html"])

    def test_non_nrw_no_nrw_norm_or_muenster_defaults(self):
        for state in ("Sachsen", "Baden-Württemberg", "", "Bayern"):
            case = case_data(0, state)
            for item in documents.build_documents(case):
                text = ParsedHTML(item["html"]).text
                self.assertNotIn("VermKatG NRW", text)
                self.assertNotIn("Münster", text)
            archive = ZipFile(BytesIO(documents.export_document(case, "zip")[0]))
            self.assertNotIn("VermKatG NRW", archive.read("quellen.txt").decode())

    def test_nrw_rule_and_conflicting_municipality_code(self):
        case = case_data()
        self.assertIn("Paragraf 14 VermKatG NRW", documents.build_documents(case)[1]["html"])
        case["selected_parcels"][0]["municipality_code"] = "14713000"
        self.assertNotIn("Paragraf 14 VermKatG NRW", documents.build_documents(case)[1]["html"])

    def test_xss_all_text_contexts_are_escaped(self):
        payload = '<img src=x onerror="alert(1)"><script>alert(2)</script>&"'
        case = case_data(1)
        for key in ("case_id", "purpose", "sender_organisation", "attachments", "legal_department_contact"):
            case[key] = payload
        case["recipients"] = dict.fromkeys(("kataster", "grundbuch", "notar"), payload)
        case["selected_parcels"][0]["location_text"] = payload
        case["profile"]["authorities"] = [{"official_name": payload}]
        for item in documents.build_documents(case):
            parsed = ParsedHTML(item["html"])
            self.assertIn(payload, parsed.text)
            self.assertNotIn("script", parsed.tags)
            self.assertNotIn("img", parsed.tags)
            self.assertFalse(any(key.startswith("on") for key, _ in parsed.attributes))
            self.assertIn("&lt;script&gt;", item["html"])
        payload_bytes = documents.export_document(case, "docx", "uebergabe")[0]
        self.assertIn(payload, "\n".join(p.text for p in Document(BytesIO(payload_bytes)).paragraphs))

    def test_complete_sentences_and_decimal_headings(self):
        for item in documents.build_documents(case_data()):
            parsed = ParsedHTML(item["html"])
            for paragraph in parsed.paragraphs:
                if paragraph in (documents.DRAFT_MARKER, "Sehr geehrte Damen und Herren,", "Mit freundlichen Grüßen"):
                    continue
                self.assertTrue(paragraph.endswith("."), paragraph)
                self.assertGreater(len(paragraph.split()), 4)
            self.assertNotRegex(item["html"], r"<h[23]>[IVXABC]+[.)]")
            self.assertNotIn("geprüft werden..", parsed.text)
            self.assertNotIn("behauptet..", parsed.text)
            self.assertNotIn("Ausgabe geprüft..", parsed.text)

    def test_alternatives_and_no_privileged_stadtwerk_claim(self):
        for item in documents.build_documents(case_data()):
            text = ParsedHTML(item["html"]).text
            self.assertIn("Grundbuchamt und Notar sind alternative Zugangswege", text)
            self.assertIn("keine automatische Zugangsberechtigung", text)
            self.assertIn("Es wurde nichts versandt und kein Notar beauftragt", text)
            self.assertIn("Grundstück im grundbuchrechtlichen Sinn", text)
            self.assertIn("mehrere Grundstücke", text)
        self.assertIn("Paragraf 133a Absatz 2 GBO", documents.build_documents(case_data())[3]["html"])

    def test_profile_statuses_remain_qualified_not_asserted(self):
        case = case_data()
        case["profile"]["authorities"] = [{"official_name": "Prüfstelle", "verification_state": "verifiziert", "verified_at": "2020-01-01"}]
        text = ParsedHTML(documents.build_documents(case)[0]["html"]).text
        self.assertIn("mitgelieferte Prüfstatus lautet verifiziert", text)
        self.assertIn("keine neue Verifikation", text)
        self.assertIn("erneut zu prüfen", text)

    def test_parcel_claims_are_qualified_without_downgrading_current_api_data(self):
        for status in ("amtliche_flurstuecksdaten", "verifiziert", "manuell_ergaenzt", "location_hint"):
            case = case_data(1)
            case["selected_parcels"][0]["identification_status"] = status
            before = deepcopy(case)
            self.assertEqual(documents.validate_case(case)["selected_parcels"], case["selected_parcels"])
            for item in documents.build_documents(case):
                text = ParsedHTML(item["html"]).text
                self.assertIn("Identifikationsstatus ausschließlich aus dem Vorgang übernommen", text)
                self.assertIn("keine serverseitige Prüfung der einzelnen Flurstücke", text)
                self.assertIn("Identifikationsstatus lautet " + status, text)
            self.assertEqual(case, before)

    def test_deterministic_and_non_mutating(self):
        case = case_data()
        before = deepcopy(case)
        self.assertEqual(documents.build_documents(case), documents.build_documents(case))
        for format in ("html", "docx", "zip", "json"):
            self.assertEqual(documents.export_document(case, format), documents.export_document(case, format))
        self.assertEqual(case, before)

    def test_docx_is_real_ooxml_with_expected_font_and_text_parity(self):
        for item in documents.build_documents(case_data()):
            data, mimetype, filename = documents.export_document(case_data(), "docx", item["id"])
            self.assertEqual(mimetype, documents.DOCX_MIMETYPE)
            self.assertTrue(filename.endswith(".docx"))
            self.assertTrue(is_zipfile(BytesIO(data)))
            archive = ZipFile(BytesIO(data))
            self.assertIn("[Content_Types].xml", archive.namelist())
            self.assertIn("word/document.xml", archive.namelist())
            document = Document(BytesIO(data))
            self.assertEqual(document.styles["Normal"].font.name, "Times New Roman")
            self.assertEqual(document.styles["Normal"].font.size.pt, 11)
            self.assertEqual(document.styles["Normal"].element.rPr.rFonts.get(qn("w:cs")), "Times New Roman")
            self.assertFalse(document.settings.odd_and_even_pages_header_footer)
            self.assertFalse(document.sections[0].different_first_page_header_footer)

            for name in ("Normal", "Title", "Heading 1", "Heading 2"):
                style = document.styles[name]
                self.assertEqual(style.font.name, "Times New Roman")
                self.assertEqual(style.font.size.pt, 11)
                self.assertIsNone(style.element.rPr.rFonts.get(qn("w:asciiTheme")))
                self.assertFalse(style.element.xpath("./w:pPr/w:pBdr"))
            paragraphs = [p.text for p in document.paragraphs]
            for paragraph in ParsedHTML(item["html"]).paragraphs:
                self.assertIn(paragraph, paragraphs)
            self.assertIn(documents.DRAFT_MARKER, document.sections[0].header.paragraphs[0].text)

    def test_combined_docx_breaks_at_draft_starts_without_empty_break_paragraphs(self):
        document = Document(BytesIO(documents.export_document(case_data(1), "docx")[0]))
        markers = [p for p in document.paragraphs if p.text == documents.DRAFT_MARKER]
        self.assertEqual(len(markers), 4)
        self.assertEqual([p.paragraph_format.page_break_before for p in markers], [False, True, True, True])
        self.assertTrue(all(p.paragraph_format.keep_with_next for p in markers))
        self.assertFalse(document.element.xpath('.//w:br[@w:type="page"]'))

    def test_zip_contains_four_docx_four_html_and_roundtrip_json(self):
        case = case_data()
        data, mimetype, filename = documents.export_document(case, "zip")
        self.assertEqual(mimetype, "application/zip")
        self.assertTrue(filename.endswith(".zip"))
        with ZipFile(BytesIO(data)) as archive:
            expected = {f"{id}.{ext}" for id in documents.TITLES for ext in ("docx", "html")}
            self.assertEqual(set(archive.namelist()), expected | {"vorgang.json", "quellen.txt"})
            self.assertEqual(json.loads(archive.read("vorgang.json")), case)
            self.assertEqual(documents.validate_case(json.loads(archive.read("vorgang.json"))), case)
            for id in documents.TITLES:
                self.assertTrue(is_zipfile(BytesIO(archive.read(id + ".docx"))))
                self.assertTrue(archive.read(id + ".html").startswith(b"<!DOCTYPE html>"))
                self.assertIn(documents.DRAFT_MARKER, archive.read(id + ".html").decode())

    def test_single_html_and_json_have_honest_extensions(self):
        for format, suffix, mime in (("html", ".html", "text/html"), ("json", ".json", "application/json")):
            data, mimetype, filename = documents.export_document(case_data(), format)
            self.assertTrue(filename.endswith(suffix))
            self.assertTrue(mimetype.startswith(mime))
            if format == "html":
                self.assertEqual(ParsedHTML(data.decode()).tags.count("article"), 4)
            else:
                self.assertEqual(json.loads(data), case_data())

    def test_unknown_formats_ids_and_path_traversal_rejected(self):
        for format in ("doc", "pdf", "../../etc/passwd", None, [], 1):
            with self.subTest(format=format), self.assertRaises(ValueError):
                documents.export_document({}, format)
        for id in ("../../etc/passwd", "unknown", [], 12):
            with self.subTest(id=id), self.assertRaises(ValueError):
                documents.export_document({}, "docx", id)
        with self.assertRaises(ValueError):
            documents.export_document({}, "zip", "grundbuch")
        case = {"case_id": "../../etc/passwd"}
        self.assertNotIn("..", documents.export_document(case, "html")[2])

    def test_no_network_actual_send_or_attachment_file_read(self):
        case = case_data()
        case["attachments"] = "/etc/passwd"
        with patch.object(socket, "socket", side_effect=AssertionError("Netzwerkzugriff")), \
                patch("builtins.open", side_effect=AssertionError("Dateizugriff")):
            documents.validate_case(case)
            documents.build_documents(case)
            documents.export_document(case, "html")
            documents.export_document(case, "json")
        # python-docx darf seine eigene feste Paketvorlage lesen, keine Vorgangspfade.
        with patch.object(socket, "socket", side_effect=AssertionError("Netzwerkzugriff")):
            documents.export_document(case, "zip")


if __name__ == "__main__":
    unittest.main()
