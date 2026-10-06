"""Deterministische Entwürfe und begrenzte, strikt typisierte Vorgangsexporte.

Dieses Modul lädt keine Quellen, liest keine Anlagenpfade und versendet nichts.
validate_case behandelt Profile als Importdaten; Laufzeitprofile verifiziert der
Server. HTML und DOCX werden aus denselben Textblöcken erzeugt.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from html import escape
from io import BytesIO
import ipaddress
import json
import math
import re
from urllib.parse import urlsplit
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo


MAX_CASE_BYTES = 2 * 1024 * 1024
MAX_PARCELS = 200
MAX_GEOMETRY_BYTES = 256 * 1024
MAX_GEOMETRY_POSITIONS = 10000
MAX_TOTAL_POSITIONS = 40000
DRAFT_MARKER = "ENTWURF - NICHT FREIGEGEBEN"
MISSING = "[ergänzen]"
DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TITLES = {
    "uebergabe": "Übergabevermerk an die Rechtsabteilung",
    "kataster": "Anfrage an die zuständige Katasterstelle",
    "grundbuch": "Antrag an das zuständige Grundbuchamt",
    "notar": "Optionaler Auftrag an einen ausgewählten Notar",
}
LEGAL_SOURCES = (
    ("Paragraf 12 GBO", "https://www.gesetze-im-internet.de/gbo/__12.html"),
    ("Paragraf 133 GBO", "https://www.gesetze-im-internet.de/gbo/__133.html"),
    ("Paragraf 133a GBO", "https://www.gesetze-im-internet.de/gbo/__133a.html"),
)
NRW_SOURCE = (
    "Paragraf 14 VermKatG NRW",
    "https://recht.nrw.de/lrgv/gesetz/08122020-gesetz-ueber-die-landesvermessung-und-das-liegenschaftskataster-vermessungs/",
)


def _fail(path: str, reason: str) -> None:
    raise ValueError(f"{path}: {reason}")


def _text(value, path):
    if type(value) is not str or len(value) > 20000:
        _fail(path, "Es ist ein Text mit höchstens 20000 Zeichen erforderlich.")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f\ud800-\udfff\ufffe\uffff]", value):
        _fail(path, "Der Text enthält unzulässige Steuerzeichen.")
    return value


def _url(value, path):
    value = _text(value, path)
    if not value:
        return value
    if len(value) > 4096 or re.search(r"[\s\\<>\"']", value):
        _fail(path, "Die Quellen-URL ist ungültig.")
    try:
        parts = urlsplit(value)
        host = parts.hostname
        port = parts.port
        if parts.scheme.lower() not in ("http", "https") or not host or parts.username or parts.password:
            raise ValueError
        if port is not None and not 1 <= port <= 65535:
            raise ValueError
        host = host.rstrip(".").lower()
        if host == "localhost" or host.endswith((".localhost", ".local", ".internal")):
            raise ValueError
        try:
            address = ipaddress.ip_address(host)
        except ValueError:
            ascii_host = host.encode("idna").decode("ascii")
            if "." not in ascii_host or not re.fullmatch(r"[a-z0-9.-]+", ascii_host):
                raise ValueError
            if all(part.isdigit() for part in ascii_host.split(".")):
                raise ValueError
        else:
            if not address.is_global:
                raise ValueError
    except (ValueError, UnicodeError):
        _fail(path, "Es ist eine öffentliche HTTP(S)-URL ohne Zugangsdaten erforderlich.")
    return value


def _number(value, path):
    if type(value) not in (int, float) or abs(value) > 1e15 or not math.isfinite(value):
        _fail(path, "Es ist eine endliche Zahl erforderlich.")
    return value


def _area(value, path):
    if value is None:
        return value
    value = _number(value, path)
    if value < 0:
        _fail(path, "Eine Fläche darf nicht negativ sein.")
    return value


def _sequence(validator, limit=200):
    def validate(value, path):
        if type(value) is not list or len(value) > limit:
            _fail(path, f"Es ist eine Liste mit höchstens {limit} Einträgen erforderlich.")
        return [validator(item, f"{path}[{index}]") for index, item in enumerate(value)]
    return validate


def _boolean(value, path):
    if type(value) is not bool:
        _fail(path, "Es ist ein Wahrheitswert erforderlich.")
    return value


def _text_or_list(value, path):
    return _sequence(_text, 100)(value, path) if type(value) is list else _text(value, path)


def _object(schema):
    def validate(value, path):
        if type(value) is not dict:
            _fail(path, "Es ist ein JSON-Objekt erforderlich.")
        if any(type(key) is not str or key not in schema for key in value):
            _fail(path, "Das Objekt enthält nicht erlaubte Felder.")
        return {key: schema[key](item, f"{path}.{key}") for key, item in value.items()}
    return validate


def _nullable(validator):
    return lambda value, path: None if value is None else validator(value, path)


def _fields(names, validator=_text):
    return {name: validator for name in names.split()}


def _position(value, path):
    coordinates = _sequence(_number, 3)(value, path)
    if len(coordinates) not in (2, 3) or not -180 <= coordinates[0] <= 180 or not -90 <= coordinates[1] <= 90:
        _fail(path, "Es werden WGS84-Koordinaten als Länge/Breite erwartet.")
    return coordinates


def _center(value, path):
    if value is None:
        return None
    if type(value) is dict:
        result = _object({"latitude": _number, "longitude": _number})(value, path)
        if set(result) != {"latitude", "longitude"}:
            _fail(path, "Breite und Länge sind erforderlich.")
        _position([result["longitude"], result["latitude"]], path)
        return result
    if type(value) is not list or len(value) != 2:
        _fail(path, "Das Kartenzentrum muss [Breite, Länge] enthalten.")
    _position([value[1], value[0]], path)
    return list(value)


def _bounds(value, path):
    if value is None:
        return None
    result = _sequence(_number, 4)(value, path)
    if len(result) != 4:
        _fail(path, "Die Begrenzung muss [West, Süd, Ost, Nord] enthalten.")
    _position(result[:2], path)
    _position(result[2:], path)
    if result[0] > result[2] or result[1] > result[3]:
        _fail(path, "Die Begrenzung ist vertauscht.")
    return result


def _geometry(value, path):
    if value is None:
        return None
    if type(value) is not dict or set(value) != {"type", "coordinates"}:
        _fail(path, "Geometrien dürfen nur type und coordinates enthalten.")
    if len(_json_bytes(value)) > MAX_GEOMETRY_BYTES:
        _fail(path, "Die Geometrie überschreitet 256 KiB.")
    count = 0

    def position(item, location):
        nonlocal count
        count += 1
        if count > MAX_GEOMETRY_POSITIONS:
            _fail(path, "Die Geometrie enthält zu viele Positionen.")
        return _position(item, location)

    def line(item, location):
        result = _sequence(position, MAX_GEOMETRY_POSITIONS)(item, location)
        if len(result) < 2:
            _fail(location, "Eine Linie braucht mindestens zwei Positionen.")
        return result

    def ring(item, location):
        result = line(item, location)
        if len(result) < 4 or result[0] != result[-1]:
            _fail(location, "Ein Polygonring braucht mindestens vier Positionen und muss geschlossen sein.")
        return result

    def nonempty(validator):
        def validate(item, location):
            result = _sequence(validator, MAX_GEOMETRY_POSITIONS)(item, location)
            if not result:
                _fail(location, "Die Geometrie darf nicht leer sein.")
            return result
        return validate

    polygon = nonempty(ring)
    validators = {"Point": position, "MultiPoint": nonempty(position), "LineString": line,
                  "MultiLineString": nonempty(line), "Polygon": polygon, "MultiPolygon": nonempty(polygon)}
    kind = value["type"]
    if type(kind) is not str or kind not in validators:
        _fail(path, "Dieser GeoJSON-Geometrietyp wird nicht unterstützt.")
    return {"type": kind, "coordinates": validators[kind](value["coordinates"], path + ".coordinates")}


EVIDENCE_SCHEMA = {
    **_fields("title name kind description note status verification_state official_name"),
    **_fields("verified_at checked_at retrieved_at source_date matched returned", _nullable(_text)),
    **_fields("url source_url tested_url", _url),
}
AUTHORITY_SCHEMA = {
    **_fields("type authority_type official_name jurisdiction postal_address visitor_address submission_information verification_state status note"),
    "verified_at": _nullable(_text),
    **_fields("official_website official_form_url source_url", _url),
    "evidence_urls": _sequence(_url, 30),
}
COVERAGE_SCHEMA = {
    **_fields("country verification_state"),
    "state_codes": _sequence(_text, 16),
    "bounds": _sequence(_bounds, 100),
}
PROVIDER_SCHEMA = {
    **_fields("id provider_id owner title role protocol version axis_order license attribution access_requirements test_result status verification_state name layer type_name reason format"),
    "verified_at": _nullable(_text),
    **_fields("catalog_source_url service_url source_url tested_url", _url),
    "layers_or_collections": _sequence(_text, 100),
    "layers": _text_or_list,
    "crs": _text_or_list,
    "coverage": _object(COVERAGE_SCHEMA),
    "requires_host_approval": _boolean,
}
PROFILE_SCHEMA = {
    **_fields("id profile_id name state state_code district municipality_code verification_state center_method attribution license"),
    **_fields("verified_at data_date", _nullable(_text)),
    **_fields("source_url district_source_url", _url),
    "source_urls": _sequence(_url, 30),
    "warnings": _sequence(_text, 100),
    "center": _center, "bounds": _bounds,
    "authorities": _sequence(_object(AUTHORITY_SCHEMA), 30),
    "providers": _sequence(_object(PROVIDER_SCHEMA), 50),
    "evidence": _sequence(_object(EVIDENCE_SCHEMA), 100),
}
PARCEL_SCHEMA = {
    **_fields("stable_id provider_id municipality municipality_code district_name district_code flur numerator denominator official_parcel_reference area_unit location_text source_date retrieved_at identification_status"),
    "source_url": _url, "area_value": _area, "geometry": _geometry,
    # Blattstellen sind optionale, manuell belegte Angaben, keine Kartenableitung.
    **_fields("grundbuchblatt grundbuchbezirk grundbuchblatt_source grundbuchblatt_date"),
    "grundbuchblaetter": _sequence(_text, 50),
}
CASE_SCHEMA = {
    **_fields("case_id purpose specific_interest requested_information_scope sender_organisation sender_address contact_person legal_department_contact attachments created_at updated_at"),
    "profile": _nullable(_object(PROFILE_SCHEMA)),
    "selected_parcels": _sequence(_object(PARCEL_SCHEMA), MAX_PARCELS),
    "recipients": _object(_fields("kataster grundbuch notar")),
}


def _json_bytes(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _preflight(value):
    nodes = 0

    def visit(item, depth=0):
        nonlocal nodes
        nodes += 1
        if depth > 16 or nodes > 150000:
            _fail("Vorgang", "Die JSON-Struktur ist zu groß oder zu tief verschachtelt.")
        if type(item) is dict:
            for key, child in item.items():
                _text(key, "Feldname")
                visit(child, depth + 1)
        elif type(item) is list:
            for child in item:
                visit(child, depth + 1)
        elif type(item) is str:
            _text(item, "Text")
        elif type(item) in (int, float):
            _number(item, "Zahl")
        elif item is not None and type(item) is not bool:
            _fail("Vorgang", "Es sind nur JSON-Datentypen erlaubt.")

    visit(value)
    size = 0
    for chunk in json.JSONEncoder(ensure_ascii=False, allow_nan=False).iterencode(value):
        size += len(chunk.encode("utf-8"))
        if size > MAX_CASE_BYTES:
            _fail("Vorgang", "Der Vorgang überschreitet 2 MiB.")


def _validate(case, *, imported):
    _preflight(case)
    result = _object(CASE_SCHEMA)(case, "Vorgang")
    positions = 0
    seen = set()
    for parcel in result.get("selected_parcels", []):
        stable_id = parcel.get("stable_id", "")
        if stable_id:
            key = (parcel.get("provider_id", ""), stable_id)
            if key in seen:
                _fail("selected_parcels", "Eine Auswahl-ID ist doppelt vorhanden.")
            seen.add(key)
        geometry = parcel.get("geometry")
        if geometry:
            pending = [geometry["coordinates"]]
            while pending:
                item = pending.pop()
                if type(item[0]) in (int, float):
                    positions += 1
                else:
                    pending.extend(item)
        if positions > MAX_TOTAL_POSITIONS:
            _fail("selected_parcels", "Die gesamte Auswahl enthält zu viele Geometriepositionen.")
    if imported and result.get("profile"):
        def downgrade(entry):
            for key in ("status", "verification_state", "test_result"):
                if entry.get(key, "").strip().casefold() in ("verifiziert", "amtlich_verifiziert", "verified", "belegt", "ok", "success", "active"):
                    entry[key] = "gefunden_ungeprueft"

        downgrade(result["profile"])
        for collection in ("providers", "authorities", "evidence"):
            for entry in result["profile"].get(collection, []):
                downgrade(entry)
    return result


def validate_case(case) -> dict:
    """Prüft ein dekodiertes JSON-Objekt und gibt eine unabhängige Kopie zurück.

    Unbekannte Felder werden abgewiesen. Fehlende Formulardaten bleiben erlaubt.
    Importierte Prüfbehauptungen sind keine aktuelle Serververifikation; sie
    werden unabhängig vom mitgelieferten Datum auf gefunden_ungeprueft gesetzt.
    URLs werden nur geprüft, nie abgerufen; ein Netzwerkadapter muss zusätzlich
    Hosts, DNS-Auflösung und Weiterleitungen selbst gegen SSRF absichern.
    """
    return _validate(case, imported=True)


def _value(mapping, key):
    value = mapping.get(key)
    return MISSING if value is None or (isinstance(value, str) and not value.strip()) else str(value)


def _sentence(text):
    return text if text.rstrip().endswith((".", "!", "?")) else text + "."


def _nrw(case):
    profile = case.get("profile") or {}
    if profile.get("state", "").strip().casefold() not in ("nordrhein-westfalen", "nrw", "de-nw", "nw"):
        return False
    codes = [profile.get("municipality_code", "")]
    codes.extend(parcel.get("municipality_code", "") for parcel in case.get("selected_parcels", []))
    return all(not code or (len(code) == 8 and code.isdigit() and code.startswith("05")) for code in codes)


def _parcel_blocks(parcels):
    blocks = [("h2", "3. Vollständige Flurstücksauswahl")]
    if parcels:
        blocks.append(("p", "Für die gesamte Auswahl werden Quellen und Identifikationsstatus ausschließlich aus dem Vorgang übernommen. "
                       "Diese Angaben bestätigen keine serverseitige Prüfung der einzelnen Flurstücke; auch eine bekannte Anbieter-ID "
                       "oder ein geprüftes Ortsprofil ist dafür kein Nachweis."))
    if not parcels:
        blocks.append(("p", "Es sind noch keine Flurstücke ausgewählt. Die Auswahl ist vor einer externen Anfrage zu ergänzen: [ergänzen]."))
    for index, parcel in enumerate(parcels, 1):
        v = lambda key: _value(parcel, key)
        parcel_number = v('numerator') + ('/' + v('denominator') if parcel.get('denominator') else '')
        blocks.extend([
            ("h3", f"3.{index}. Auswahlposition {index}"),
            ("p", f"Die Gemeinde ist mit {v('municipality')} und dem Gemeindeschlüssel {v('municipality_code')} angegeben. "
             f"Die Gemarkung heißt {v('district_name')}; ihr Schlüssel lautet {v('district_code')}. "
             f"Die Flur ist mit {v('flur')} und die Flurstücksnummer mit {parcel_number} angegeben. "
             f"Das amtliche Flurstückskennzeichen lautet nach den Vorgangsdaten {v('official_parcel_reference')}."),
            ("p", f"Die angegebene Fläche beträgt {v('area_value')} {v('area_unit')}. Als Lage ist festgehalten: {v('location_text')}. "
             f"Die Quelle ist {v('source_url')}. Der Datenstand lautet {v('source_date')}; der Abrufzeitpunkt lautet {v('retrieved_at')}. "
             f"Der eingetragene Identifikationsstatus lautet {v('identification_status')}. "
             f"Die interne Auswahl-ID lautet {v('stable_id')}; die Anbieter-ID lautet {v('provider_id')}."),
        ])
        status = parcel.get("identification_status", "").casefold()
        if "manuell" in status or "manual" in status:
            note = "Diese Auswahl wurde manuell ergänzt. Ein manueller Nachweis ist gesondert zu prüfen; die Eingabe ist kein amtlicher Eigentümernachweis."
        elif "lage" in status or "location_hint" in status or not parcel.get("official_parcel_reference"):
            note = "Diese Auswahl ist zunächst ein Lagehinweis oder unvollständig identifiziert. Eine amtliche Flurstücksidentifikation ist einzuholen; Nummern werden nicht aus der Geometrie abgeleitet."
        else:
            note = "Die Identifikationsangaben werden aus dem Vorgang übernommen und sind anhand der angegebenen Quelle zu prüfen. Auch ein als amtlich bezeichnetes Kartenmerkmal ist kein Eigentümernachweis."
        blocks.append(("p", note))
        sheets = list(parcel.get("grundbuchblaetter", []))
        if parcel.get("grundbuchblatt"):
            sheets.insert(0, parcel["grundbuchblatt"])
        sheet_text = "; ".join(sheets) if sheets else MISSING
        blocks.append(("p", f"Als Grundbuchblatt oder Blattstellen sind {sheet_text} im Grundbuchbezirk {v('grundbuchbezirk')} angegeben. "
                       f"Diese manuell beigebrachten Angaben sind anhand des Nachweises {v('grundbuchblatt_source')} "
                       f"mit dem Nachweisdatum {v('grundbuchblatt_date')} zu prüfen. Eine fehlende oder mehrfache Zuordnung bleibt offen; "
                       "es wird keine Blattnummer erfunden und keine eindeutige Zuordnung aus der Karte behauptet."))
    return blocks


def _authority_blocks(case):
    blocks = [("h2", "5. Empfänger, Quellen und Prüfstand")]
    profile = case.get("profile") or {}
    blocks.append(("p", f"Die Verwaltungszuordnung des Ortsprofils verweist auf {_value(profile, 'source_url')}. "
                   f"Der überlieferte Prüfstatus lautet {_value(profile, 'verification_state')}; das Prüfdatum lautet "
                   f"{_value(profile, 'verified_at')} und der Datenstand lautet {_value(profile, 'data_date')}. "
                   f"Als Kreis oder kreisfreie Stadt ist {_value(profile, 'district')} mit der Quelle "
                   f"{_value(profile, 'district_source_url')} angegeben."))
    for warning in profile.get("warnings", []):
        blocks.append(("p", _sentence(f"Als offener Prüfpunkt ist im Ortsprofil vermerkt: {warning}")))
    recipients = case.get("recipients", {})
    for key, name in (("kataster", "Katasterstelle"), ("grundbuch", "Grundbuchamt"), ("notar", "Notar")):
        blocks.append(("p", f"Für den Weg über {name} ist folgender Empfänger eingetragen: {_value(recipients, key)}. "
                       "Diese Texteingabe ersetzt keine Prüfung der Zuständigkeit, Auswahl oder postalischen Anschrift."))
    for authority in profile.get("authorities", []):
        v = lambda key: _value(authority, key)
        urls = "; ".join(authority.get("evidence_urls", [])) or v("source_url")
        blocks.append(("p", f"Das Ortsprofil nennt {v('official_name')} als Stelle des Typs {authority.get('type') or v('authority_type')} "
                       f"für den Zuständigkeitsbereich {v('jurisdiction')}. Die Briefanschrift lautet {v('postal_address')}; "
                       f"die davon getrennte Besucheranschrift lautet {v('visitor_address')}. "
                       f"Der mitgelieferte Prüfstatus lautet {v('verification_state')} und der weitere Status lautet {v('status')}; "
                       f"das überlieferte Prüfdatum lautet {v('verified_at')}. "
                       f"Die Nachweisquellen sind {urls}; die Verfahrensseite ist {v('official_website')} "
                       f"und der Formularlink ist {v('official_form_url')}. "
                       f"Zum Einreichungsweg ist angegeben: {v('submission_information')}."))
    if not profile.get("authorities"):
        blocks.append(("p", "Eine belegte behördliche Zuständigkeit ist im Ortsprofil nicht hinterlegt und muss geprüft werden: [ergänzen]."))
    for provider in profile.get("providers", []):
        blocks.append(("p", f"Für den Anbieter {_value(provider, 'provider_id')} nennt das Profil die Quelle "
                       f"{_value(provider, 'catalog_source_url')} und den Dienst {_value(provider, 'service_url')}. "
                       f"Der mitgelieferte Prüfstatus lautet {_value(provider, 'verification_state')} und der weitere Status lautet "
                       f"{_value(provider, 'status')}; der Prüfvermerk lautet {_value(provider, 'test_result')} mit dem Prüfdatum "
                       f"{_value(provider, 'verified_at')}. Als Lizenz ist {_value(provider, 'license')} angegeben; "
                       f"der Herkunftsnachweis lautet {_value(provider, 'attribution')}."))
    for evidence in profile.get("evidence", []):
        blocks.append(("p", f"Der Profilnachweis {_value(evidence, 'title')} nennt als Quelle "
                       f"{evidence.get('source_url') or evidence.get('url') or MISSING}. "
                       f"Der mitgelieferte Prüfstatus lautet {_value(evidence, 'verification_state')} und der weitere Status lautet {_value(evidence, 'status')}; "
                       f"das Prüfdatum lautet {evidence.get('verified_at') or evidence.get('checked_at') or MISSING} "
                       f"und der Abrufzeitpunkt lautet {_value(evidence, 'retrieved_at')}."))
    blocks.append(("p", "Ein mitgelieferter Prüfstatus dokumentiert lediglich den Vorgangsstand und ist keine neue Verifikation. "
                   "Importierte oder veraltete Profile sind erneut zu prüfen. Eine Kontakt-E-Mail belegt keinen zulässigen Einreichungsweg."))
    return blocks


@dataclass(frozen=True)
class _Draft:
    id: str
    title: str
    blocks: tuple[tuple[str, str], ...]


def _drafts(case):
    v = lambda key: _value(case, key)
    profile = case.get("profile") or {}
    parcels = case.get("selected_parcels", [])
    common = [
        ("p", "Dieses Dokument ist ein nicht freigegebener Entwurf. Es wurde nichts versandt und kein Notar beauftragt."),
        ("h2", "1. Vorgang und Beteiligte"),
        ("p", f"Der Vorgang trägt die Kennung {v('case_id')}. Er wurde am {v('created_at')} angelegt und zuletzt am {v('updated_at')} geändert. "
         f"Das Ortsprofil nennt {_value(profile, 'name')} im Bundesland {_value(profile, 'state')} "
         f"mit dem Gemeindeschlüssel {_value(profile, 'municipality_code')}."),
        ("p", f"Als Absender ist {v('sender_organisation')} mit der Anschrift {v('sender_address')} vorgesehen. "
         f"Ansprechperson ist {v('contact_person')}. Für die rechtliche Prüfung ist {v('legal_department_contact')} benannt."),
        ("h2", "2. Anlass und konkretes Interesse"),
        ("p", " ".join((
            _sentence(f"Für die Recherche ist folgender Zweck angegeben: {v('purpose')}"),
            _sentence(f"Das konkrete Interesse wird im Vorgang wie folgt begründet: {v('specific_interest')}"),
            _sentence(f"Der gewünschte Auskunftsumfang ist wie folgt beschrieben: {v('requested_information_scope')}"),
        ))),
    ]
    definitions = ("p", "Ein Flurstück ist eine im Liegenschaftskataster abgegrenzte und bezeichnete Bodenfläche. "
                   "Ein Grundstück im grundbuchrechtlichen Sinn ist die unter einer eigenen laufenden Nummer im Bestandsverzeichnis gebuchte rechtliche Einheit; "
                   "es kann mehrere Flurstücke umfassen. Ein Grundbuchblatt ist das Registerblatt und kann mehrere Grundstücke enthalten. "
                   "Flurstückskennzeichen und Grundbuchblattnummer sind nicht austauschbar. Offene Karten und Geometrien weisen kein Eigentum nach.")
    nrw = ("Für die Katasterauskunft in Nordrhein-Westfalen ist Paragraf 14 VermKatG NRW zu prüfen. "
           "Eigentümerangaben setzen nach Absatz 2 grundsätzlich die Darlegung eines berechtigten Interesses voraus. "
           "Die gesetzlichen Ausnahmen sind anhand der konkreten Rolle zu prüfen. Nach Absatz 3 sind die Eigentümerangaben nach Zweckerfüllung zu löschen; "
           "ein Datenbestand für unbestimmte Zwecke ist unzulässig.") if _nrw(case) else (
           f"Die landesrechtliche Grundlage für Kataster- und Eigentümerauskünfte im angegebenen Bundesland {_value(profile, 'state')} "
           "ist noch anhand der zuständigen amtlichen Quellen zu prüfen: [ergänzen]. Eine landesspezifische Zugangsberechtigung wird nicht vorausgesetzt.")
    routes = {
        "uebergabe": [
            "Die Rechtsabteilung wird gebeten, den Vorgang sowie das konkrete Interesse, die Vertretungsbefugnis und den erforderlichen Auskunftsumfang zu prüfen. "
            "Sind Identität oder Zuordnung der ausgewählten Flächen unklar, ist zunächst die zuständige Katasterstelle um Identifizierung zu bitten. "
            "Erst nach dieser Klärung ist über eine zulässige Eigentümerauskunft oder Grundbucheinsicht zu entscheiden.",
            nrw,
            "Paragraf 12 Absatz 1 GBO verlangt die Darlegung eines berechtigten Interesses; Absatz 2 betrifft Abschriften im zulässigen Einsichtsumfang. "
            "Paragraf 133 GBO regelt ein gesondertes, genehmigungsbedürftiges automatisiertes Abrufverfahren und schafft hier keinen Zugang. "
            "Paragraf 133a GBO betrifft notarielle Mitteilungen bei dargelegtem berechtigtem Interesse; öffentliche Interessen allein genügen für diesen Weg nicht.",
        ],
        "kataster": [
            "Sehr geehrte Damen und Herren,",
            "wir bitten um Prüfung unserer Anfrage zu den unter Nummer 3 vollständig aufgeführten Flächen. "
            "Soweit die Flurstücksidentifikation noch fehlt oder nur manuelle Lagehinweise vorliegen, bitten wir zunächst um Identifizierung und Zuordnung. "
            "Für die eindeutig identifizierten Flurstücke bitten wir im oben begründeten Umfang um Katasterauskunft und, soweit rechtlich zulässig und erforderlich, um Eigentümerangaben.",
            nrw,
            "Bitte teilen Sie uns mit, wenn die Darlegung unseres Interesses, die Vertretungsnachweise oder die Angaben zum Auskunftsumfang ergänzt werden müssen. "
            "Wir bitten vor einer kostenpflichtigen Bearbeitung um Mitteilung der voraussichtlichen Kosten und etwaiger verbindlicher Formulare.",
        ],
        "grundbuch": [
            "Sehr geehrte Damen und Herren,",
            "wir beantragen nach Paragraf 12 GBO Zugang zu den Grundbuchinformationen, die dem oben beschriebenen Zweck und konkreten Interesse entsprechen. "
            + _sentence(f"Für Art und Umfang des Antrags ist folgende Angabe maßgeblich: {v('requested_information_scope')}") + " "
            "Eine Einsicht, einfache Abschrift beziehungsweise ein einfacher Ausdruck oder eine beglaubigte Abschrift beziehungsweise ein amtlicher Ausdruck "
            "ist vor Einreichung eindeutig auszuwählen, soweit dies aus dieser Angabe noch nicht hervorgeht: [ergänzen]. "
            "Der Antrag erstreckt sich nicht pauschal auf sämtliche Abteilungen oder Grundakten.",
            "Die beigefügten Flurstücksangaben dienen der Zuordnung und ersetzen keine bekannte Grundbuchblattstelle. "
            "Soweit die Blattstelle unbekannt oder mehrdeutig ist, bitten wir zunächst um Prüfung der Zuordnung im Rahmen des zulässigen Zugangs. "
            "Bitte weisen Sie uns auf erforderliche weitere Belege zum Interesse und zur Vertretungsbefugnis hin. "
            "Paragraf 133 GBO betrifft ein gesondertes automatisiertes Abrufverfahren; dessen Zulassung wird mit diesem Antrag weder behauptet noch ersetzt.",
        ],
        "notar": [
            "Sehr geehrte Damen und Herren,",
            "für den Fall der gesonderten Freigabe durch unsere Rechtsabteilung bitten wir Sie um Prüfung, ob und in welchem Umfang "
            "Sie uns die benötigten Grundbuchinformationen nach Paragraf 133a Absatz 1 GBO mitteilen dürfen. "
            "Das hierfür erforderliche berechtigte Interesse im Sinne von Paragraf 12 GBO ist anhand der vorstehenden Angaben und der genannten Unterlagen zu prüfen. "
            "Ein allgemeines öffentliches Interesse ersetzt diese Darlegung nicht; Paragraf 133a Absatz 2 GBO schließt Mitteilungen im öffentlichen Interesse "
            "oder zu wissenschaftlichen und Forschungszwecken aus. Etwaige landesrechtliche Einschränkungen nach Absatz 5 sind ebenfalls zu prüfen.",
            "Bitte klären Sie vor der kostenpflichtigen Tätigkeit den voraussichtlichen Kostenrahmen und gegebenenfalls weitere notwendige Nachweise. "
            "Die Organisation wählt die Notarin oder den Notar selbst aus. Dieser Entwurf enthält keine bereits erteilte Beauftragung "
            "und begründet keinen Anspruch auf eine Mitteilung oder einen automatisierten Abruf.",
        ],
    }
    final = [
        ("h2", "6. Unterlagen und Freigabe"),
        ("p", _sentence(f"Im Vorgang sind folgende Anlagen oder Nachweise benannt: {v('attachments')}") + " "
         "Die Benennung bestätigt weder deren Beifügung noch deren Prüfung. Fehlende Nachweise zum konkreten Interesse und zur Vertretungsbefugnis sind zu ergänzen: [ergänzen]."),
        ("p", "Stadtwerke oder Wasserwerke besitzen allein aufgrund ihrer Versorgungsaufgabe keine automatische Zugangsberechtigung "
         "und werden nicht pauschal als Behörden mit Sonderzugriff behandelt. Über den Zugang entscheidet die zuständige Stelle anhand des Einzelfalls."),
        ("p", "Grundbuchamt und Notar sind alternative Zugangswege. Die Entwürfe werden nicht automatisch parallel versandt. "
         "Die Rechtsabteilung prüft Zuständigkeit, Interesse, Vertretung, Umfang, Kosten und Einreichungsweg und entscheidet über die Freigabe. "
         "Das Tool versendet keine Anfragen und erteilt keine Aufträge."),
        ("p", "Die Freigabeentscheidung, die vertretungsberechtigte unterzeichnende Person und das Datum sind vor Verwendung zu ergänzen: [ergänzen]."),
    ]
    result = []
    for document_id, title in TITLES.items():
        recipient = v("legal_department_contact") if document_id == "uebergabe" else _value(case.get("recipients", {}), document_id)
        blocks = common + [("p", f"Als Empfänger dieses Entwurfs ist vorgesehen: {recipient}.")]
        blocks += _parcel_blocks(parcels) + [definitions, ("h2", "4. Prüfbitte und Auskunftsweg")]
        blocks += [("p", paragraph) for paragraph in routes[document_id]]
        if document_id == "uebergabe":
            blocks += _authority_blocks(case)
        else:
            blocks += [("h2", "5. Zuständigkeit und Einreichung"),
                       ("p", "Der eingetragene Empfänger und seine Anschrift sind vor Einreichung anhand der amtlichen Zuständigkeits- und Verfahrensangaben zu prüfen. "
                        "Ein kommunaler Gebietsname allein belegt keine Grundbuchamtszuständigkeit. Eine Kontakt-E-Mail ist nicht ohne Weiteres ein zulässiger Einreichungsweg.")]
        blocks += final
        if document_id != "uebergabe":
            blocks += [("p", "Mit freundlichen Grüßen"), ("p", f"Für {v('sender_organisation')} zeichnet nach Freigabe die vertretungsberechtigte Person: [ergänzen].")]
        result.append(_Draft(document_id, title, tuple(blocks)))
    return result


def _html(drafts):
    articles = []
    for draft in drafts:
        body = "\n".join(f"<{tag}>{escape(text, quote=True)}</{tag}>" for tag, text in draft.blocks)
        articles.append(f'<article id="{draft.id}"><p class="draft">{DRAFT_MARKER}</p>'
                        f'<h1>{escape(draft.title)}</h1>\n{body}</article>')
    title = drafts[0].title if len(drafts) == 1 else "Grundstücksrecherche: vier Entwürfe"
    return ('<!DOCTYPE html>\n<html lang="de"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width, initial-scale=1">'
            '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; base-uri \'none\'; form-action \'none\'">'
            f'<title>{escape(title)}</title><style>'
            'body{font-family:"Times New Roman",Times,serif;font-size:11pt;line-height:1.3;'
            'max-width:170mm;margin:12mm auto;padding:0 6mm;color:#000;background:#fff;letter-spacing:0}'
            'h1,h2,h3{font-family:inherit;font-size:11pt;color:#000;break-after:avoid;margin:16pt 0 11pt}'
            'p{margin:0 0 8pt;white-space:pre-wrap;overflow-wrap:anywhere}'
            '.draft{font-weight:bold}article+article{break-before:page;margin-top:24pt}'
            '@page{size:A4;margin:20mm}@media print{body{margin:0;padding:0;max-width:none}}'
            '</style></head><body>\n' + "\n".join(articles) + '\n</body></html>')


def build_documents(case) -> list[dict]:
    """Erzeugt stets vier vollständige HTML-Dokumente ohne Seiteneffekte."""
    checked = _validate(case, imported=False)
    return [{"id": draft.id, "title": draft.title, "html": _html([draft])} for draft in _drafts(checked)]


def _zip(entries):
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in entries:
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            archive.writestr(info, content)
    return buffer.getvalue()


def _docx(drafts):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    document = Document()
    document.settings.odd_and_even_pages_header_footer = False
    section = document.sections[0]
    section.different_first_page_header_footer = False
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{key}"), None)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), "Times New Roman")
        for child in list(style.element.rPr):
            if child.tag == qn("w:spacing"):
                style.element.rPr.remove(child)
            if child.tag == qn("w:szCs"):
                child.set(qn("w:val"), "22")
        if style.element.pPr is not None:
            for child in list(style.element.pPr):
                if child.tag in (qn("w:pBdr"), qn("w:contextualSpacing")):
                    style.element.pPr.remove(child)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.12
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].paragraph_format.space_before = Pt(12)
        document.styles[name].paragraph_format.space_after = Pt(11)
        document.styles[name].font.bold = True
    document.styles["Title"].font.bold = True
    section.header.paragraphs[0].text = DRAFT_MARKER
    for index, draft in enumerate(drafts):
        marker = document.add_paragraph(DRAFT_MARKER)
        marker.paragraph_format.page_break_before = bool(index)
        marker.paragraph_format.keep_with_next = True
        marker.runs[0].bold = True
        document.add_paragraph(draft.title, "Title")
        for kind, text in draft.blocks:
            document.add_paragraph(text, {"h2": "Heading 1", "h3": "Heading 2"}.get(kind, "Normal"))
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = drafts[0].title if len(drafts) == 1 else "Grundstücksrecherche: vier Entwürfe"
    properties.subject = DRAFT_MARKER
    properties.comments = ""
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    stream = BytesIO()
    document.save(stream)
    # python-docx vergibt ZIP-Zeitstempel; feste Zeitstempel machen Exporte reproduzierbar.
    with ZipFile(BytesIO(stream.getvalue())) as archive:
        return _zip((name, archive.read(name)) for name in sorted(archive.namelist()))


def _sources(case):
    lines = ["Öffentliche Rechtsquellen; geprüft am 06.10.2026. Vor Einreichung erneut prüfen."]
    for title, url in (*LEGAL_SOURCES, *((NRW_SOURCE,) if _nrw(case) else ())):
        lines.append(f"{title}: {url}")
    lines.append("Vorgangsquellen sind übernommene Angaben, keine neue Verifikation.")

    def collect(value):
        if type(value) is dict:
            for key, item in value.items():
                if key.endswith("url") or key in ("official_website", "evidence_urls", "source_urls"):
                    for url in item if type(item) is list else [item]:
                        if url and url not in lines:
                            lines.append(url)
                elif type(item) in (dict, list):
                    collect(item)
        elif type(value) is list:
            for item in value:
                collect(item)
    collect(case)
    return ("\n".join(lines) + "\n").encode("utf-8")


def export_document(case, format, document_id=None) -> tuple[bytes, str, str]:
    """Exportiert Bytes, MIME-Typ und sicheren Dateinamen, niemals einen Dateipfad.

    html/docx ohne ID enthalten alle vier Entwürfe. zip enthält immer vier
    einzelne DOCX, vier HTML, vorgang.json und quellen.txt. json ist der alleinige
    Vorgangsexport. Andere Formate (insbesondere als .doc umbenanntes HTML) sind
    nicht zulässig. PDF wird über die Druckansicht der vollständigen HTML erzeugt.
    """
    if type(format) is not str or format not in ("html", "docx", "zip", "json"):
        raise ValueError("Das Exportformat muss html, docx, zip oder json sein.")
    if document_id is not None and (type(document_id) is not str or document_id not in TITLES):
        raise ValueError("Die Dokument-ID ist unbekannt.")
    if format in ("zip", "json") and document_id is not None:
        raise ValueError("Gesamtexporte dürfen keine einzelne Dokument-ID angeben.")
    checked = _validate(case, imported=False)
    drafts = _drafts(checked)
    if document_id:
        drafts = [draft for draft in drafts if draft.id == document_id]
    filename = document_id or "grundstuecksrecherche-entwuerfe"
    if format == "html":
        return _html(drafts).encode("utf-8"), "text/html; charset=utf-8", filename + ".html"
    if format == "docx":
        return _docx(drafts), DOCX_MIMETYPE, filename + ".docx"
    if format == "json":
        return _json_bytes(checked), "application/json; charset=utf-8", "vorgang.json"
    entries = []
    for draft in drafts:
        entries.extend([(draft.id + ".docx", _docx([draft])),
                        (draft.id + ".html", _html([draft]).encode("utf-8"))])
    entries.extend([("vorgang.json", _json_bytes(checked)), ("quellen.txt", _sources(checked))])
    return _zip(entries), "application/zip", filename + ".zip"
