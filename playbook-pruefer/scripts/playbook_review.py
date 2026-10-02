#!/usr/bin/env python3
"""Validate evidence and aggregate supplied legal observations; never infer them.

Python 3.10+. Core: standard library. Optional PDF input: pypdf;
optional Word export: python-docx. No network requests or document execution.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from copy import deepcopy
from email import policy
from email.parser import BytesParser
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

OUTCOMES = {"met", "not_met", "not_verifiable", "pending", "excluded"}
LABELS = {"met": "Met", "not_met": "Not met", "not_verifiable": "Not verifiable", "pending": "Pending", "excluded": "Excluded"}
BADGES = {"none": "No identified playbook risk", "medium": "Medium risk", "high": "High risk", "not_verifiable": "Not verifiable"}
POSITION_LABELS = {"starting": "Ausgangsposition", "fallback": "Rückfallposition", "not_acceptable": "Rote Linie"}
MATCH_LABELS = {"true": "erfüllt", "false": "nicht erfüllt", "unknown": "ungeklärt"}


class ReviewError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise ReviewError(message)


def string(value, label):
    require(isinstance(value, str) and bool(value.strip()), f"{label}: nicht leere Zeichenfolge erforderlich")
    return value


def unique_objects(items, label):
    require(isinstance(items, list) and bool(items), f"{label}: nicht leere Liste erforderlich")
    result = {}
    for item in items:
        require(isinstance(item, dict), f"{label}: Objekt erforderlich")
        key = string(item.get("id"), f"{label}.id")
        require(key not in result, f"{label}: doppelte ID {key}")
        result[key] = item
    return result


def normalize(text):
    # Whitespace is normalized; wording, spelling and punctuation are not.
    return " ".join(text.split())


def source_file(root, relative):
    string(relative, "Dokumentpfad")
    p = Path(relative)
    require(not p.is_absolute() and ".." not in p.parts and "\\" not in relative, "Dokumentpfad muss relativ innerhalb des Quellenordners liegen")
    resolved = (root / p).resolve()
    require(resolved.is_relative_to(root.resolve()), "Dokumentpfad verlässt Quellenordner")
    require(resolved.is_file() and resolved.stat().st_size <= 50_000_000, "Dokument fehlt oder ist größer als 50 MB")
    return resolved


def extract(path):
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    if ext == ".docx":
        with ZipFile(path) as archive:
            info = archive.getinfo("word/document.xml")
            require(info.file_size <= 20_000_000, "Word-Haupttext ist zu groß")
            body = ET.fromstring(archive.read(info))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        require(not body.findall(".//w:ins", ns) and not body.findall(".//w:del", ns), "Word enthält offene Änderungen: maßgebliche Fassung erst verbindlich klären")
        return "\n".join("".join(p.itertext()) for p in body.findall(".//w:p", ns))
    if ext == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        return "\n".join(p.get_content() for p in message.walk() if p.get_content_type() == "text/plain" and p.get_content_disposition() != "attachment")
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ReviewError("PDF-Eingabe benötigt pypdf; keine automatische Installation") from exc
        reader = PdfReader(path)
        require(not reader.is_encrypted, "Verschlüsselte PDF zuerst freigeben")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    raise ReviewError(f"Nicht unterstütztes Quellenformat: {ext}")


def match(values, mode):
    values = [v for v in values if v != "excluded"]
    if not values:
        return "unknown"
    if mode == "all":
        return "false" if "not_met" in values else "true" if all(v == "met" for v in values) else "unknown"
    return "true" if "met" in values else "false" if all(v == "not_met" for v in values) else "unknown"


def validate_observation(obs, documents, authoritative, root_texts, final):
    outcome = obs.get("outcome")
    require(outcome in OUTCOMES, "Unbekannter Regelstatus")
    string(obs.get("reasoning"), "Regelbegründung")
    require(not final or outcome != "pending", "Abschlusslauf enthält Pending")
    if outcome == "excluded":
        scope = obs.get("scope_approval", {})
        string(scope.get("by"), "Ausnahmefreigabe durch")
        string(scope.get("reason"), "Begründung der sachlichen Nichtanwendbarkeit")
        string(scope.get("reference"), "Beleg der vorab erteilten Scope-Freigabe")
        require(not obs.get("citations"), "Ausgeschlossene Regel nicht als Vertragsbefund ausgeben")
        return
    citations = obs.get("citations", [])
    require(isinstance(citations, list), "Fundstellen müssen eine Liste sein")
    operative_quotes = 0
    for citation in citations:
        require(isinstance(citation, dict), "Fundstelle muss ein Objekt sein")
        did = citation.get("document_id")
        require(did in documents, "Fundstelle verweist auf unbekanntes Dokument")
        require(documents[did]["role"] != "superseded", "Überholte Fassung darf keinen aktuellen Regelbefund tragen")
        string(citation.get("locator"), "Fundstellenort")
        quote = normalize(string(citation.get("quote"), "Wörtliches Zitat"))
        require(quote in normalize(root_texts[did]), f"Zitat fehlt im Original {did}")
        require(citation.get("locator_verified") is True, "Seiten-/Klauselort muss separat am Original geprüft sein")
        if did in authoritative:
            operative_quotes += 1
    basis = obs.get("basis")
    require(basis in {"quote", "absence", "unresolved"}, "Regelbasis muss quote/absence/unresolved sein")
    if outcome in {"met", "not_met"}:
        require(basis in {"quote", "absence"}, "Entschiedene Regel ohne tragfähige Grundlage")
        if basis == "quote":
            require(operative_quotes > 0, "Entschiedene Vertragsregel benötigt mindestens ein Zitat der maßgeblichen Fassung")
        else:
            scope = obs.get("searched_document_ids")
            require(isinstance(scope, list) and len(scope) == len(set(scope)) and set(scope) == authoritative,
                    "Fehlstellenbefund benötigt vollständigen maßgeblichen Suchumfang")
            string(obs.get("search_description"), "Inhalt und Methode der Fehlstellensuche")
    else:
        require(basis == "unresolved", "Offener Regelstatus benötigt unresolved und eine konkrete Rückfrage")
        string(obs.get("question"), "Nächste Rückfrage")


def evaluate(data, source_root, *, final=True):
    require(isinstance(data, dict) and data.get("schema_version") == 1, "Unbekannte Formatversion")
    pb = data.get("playbook", {})
    for key in ("id", "version", "approved_by", "approval_reference", "contract_type", "client_side"):
        string(pb.get(key), f"Playbook.{key}")
    string(data.get("review_id"), "Review-ID")
    string(data.get("document_precedence"), "Vereinbarte Dokumentrangfolge")
    documents = unique_objects(data.get("documents"), "Dokumente")
    texts, authoritative = {}, set()
    for did, doc in documents.items():
        for key in ("title", "version"):
            string(doc.get(key), f"Dokument.{key}")
        require(doc.get("role") in {"authoritative", "context", "superseded"}, "Unbekannte Dokumentrolle")
        path = source_file(Path(source_root), doc.get("path"))
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        require(doc.get("sha256") == sha, f"Quellendatei geändert oder Hash fehlt: {did}")
        texts[did] = extract(path)
        if doc["role"] == "authoritative":
            authoritative.add(did)
            require(bool(normalize(texts[did])), f"Maßgeblicher Text nicht lesbar: {did}; OCR/Original klären")
    require(bool(authoritative), "Kein maßgebliches Vertragsdokument")
    topics = unique_objects(pb.get("topics"), "Topics")
    observations = data.get("observations")
    require(isinstance(observations, dict), "Regelbefunde als ID-Objekt erforderlich")
    supplied_topics = data.get("topic_observations", {})
    require(set(supplied_topics) == set(topics), "Topic-Befunde fehlen oder sind fremd")
    seen_rules, seen_positions, result_topics = set(), set(), []
    for tid, topic in topics.items():
        string(topic.get("title"), "Topic-Titel")
        require(type(topic.get("required")) is bool, "required muss explizit true/false sein")
        status = supplied_topics[tid]
        require(status.get("presence") in {"found", "not_found", "not_verifiable"}, "Unbekannter Topic-Fundstatus")
        string(status.get("reasoning"), "Topic-Begründung")
        positions = unique_objects(topic.get("positions"), f"Positionen {tid}")
        require(sum(p.get("type") == "starting" for p in positions.values()) == 1, "Topic benötigt genau eine Startposition")
        require(sum(p.get("type") == "not_acceptable" for p in positions.values()) >= 1, "Topic benötigt eine rote Position")
        pos_results = []
        unresolved = False
        for pid, position in positions.items():
            require(pid not in seen_positions, "Positions-ID muss global eindeutig sein")
            seen_positions.add(pid)
            kind = position.get("type")
            require(kind in {"starting", "fallback", "not_acceptable"}, "Unbekannter Positionstyp")
            mode = position.get("match_mode", "all" if kind != "not_acceptable" else None)
            require(mode in {"all", "any"}, "Positionslogik all/any muss für rote Positionen ausdrücklich feststehen")
            require(kind == "not_acceptable" or mode == "all", "Annahmefähige Positionen müssen alle ihre Regeln erfüllen")
            rules = unique_objects(position.get("rules"), f"Regeln {pid}")
            counts = {key: 0 for key in OUTCOMES}
            rows = []
            for rid, rule in rules.items():
                require(rid not in seen_rules, "Regel-ID muss global eindeutig sein")
                seen_rules.add(rid)
                string(rule.get("condition"), "Testbare Regelbedingung")
                require(rid in observations, f"Regelbefund fehlt: {rid}")
                obs = observations[rid]
                require(isinstance(obs, dict), "Regelbefund muss ein Objekt sein")
                validate_observation(obs, documents, authoritative, texts, final)
                out = obs["outcome"]
                if status["presence"] == "not_found" and out in {"met", "not_met"}:
                    require(obs["basis"] == "absence", "Topic Not found kollidiert mit einem positiven Textfund")
                counts[out] += 1
                unresolved |= out in {"not_verifiable", "pending"}
                label = ("Detected" if out == "met" else "Not detected") if kind == "not_acceptable" and out in {"met", "not_met"} else LABELS[out]
                rows.append({"id": rid, "condition": rule["condition"], **deepcopy(obs), "label": label})
            pos_results.append({"id": pid, "type": kind, "match_mode": mode,
                                "match": match([r["outcome"] for r in rows], mode),
                                "met": counts["met"], "judged": counts["met"] + counts["not_met"],
                                "not_verifiable": counts["not_verifiable"], "pending": counts["pending"],
                                "excluded": counts["excluded"], "rules": rows})
        findings = status.get("legal_findings", [])
        require(isinstance(findings, list), "Rechtsbefunde müssen eine Liste sein")
        for finding in findings:
            require(finding.get("status") in {"confirmed_invalid", "needs_review"}, "Unbekannter Rechtsbefundstatus")
            for key in ("reasoning", "authority", "source_url", "pinpoint", "checked_on"):
                string(finding.get(key), f"Rechtsbefund.{key}")
            require(finding["source_url"].startswith("https://"), "Rechtsquelle benötigt HTTPS-Link")
            require(bool(finding.get("contract_citations")), "Rechtsbefund benötigt konkrete Vertragsfundstelle")
            validate_observation({"outcome": "met", "basis": "quote", "reasoning": finding["reasoning"],
                                  "citations": finding["contract_citations"]}, documents, authoritative, texts, final)
        red = [p["match"] for p in pos_results if p["type"] == "not_acceptable"]
        acceptable = [p for p in pos_results if p["type"] != "not_acceptable"]
        satisfied = next((p for p in acceptable if p["type"] == "starting" and p["match"] == "true"), None)
        if not satisfied:
            satisfied = next((p for p in acceptable if p["match"] == "true"), None)
        no_acceptable_position = status["presence"] == "found" and all(p["match"] == "false" for p in acceptable)
        if "true" in red or any(f["status"] == "confirmed_invalid" for f in findings) or (topic["required"] and status["presence"] == "not_found") or no_acceptable_position:
            risk = "high"
        elif status["presence"] != "found" or "unknown" in red or findings:
            risk = "not_verifiable"
        elif satisfied:
            risk = "none" if satisfied["type"] == "starting" else "medium"
        else:
            risk = "not_verifiable"
        result_topics.append({"id": tid, "title": topic["title"], "required": topic["required"],
                              "presence": status["presence"], "risk": risk,
                              "badge": "Not found" if status["presence"] == "not_found" else BADGES[risk],
                              "matched_position": satisfied["id"] if satisfied else None,
                              "reasoning": status["reasoning"], "unresolved_rules": unresolved,
                              "legal_findings": deepcopy(findings), "positions": pos_results})
    require(set(observations) == seen_rules, "Fremde Regelbefunde außerhalb des Playbooks")
    return {"schema_version": 1, "review_id": data["review_id"], "mode": "final" if final else "draft",
            "playbook": {k: v for k, v in pb.items() if k != "topics"},
            "document_precedence": data["document_precedence"], "documents": deepcopy(list(documents.values())),
            "topics": result_topics,
            "ready_for_human_decision": final and all(t["risk"] in {"none", "medium"} and not t["unresolved_rules"] for t in result_topics),
            "limits": "Keine automatische Vertragsfreigabe. Zitattext maschinell geprüft; Fundstellenort, Auslegung, Rechtsbefund und Vollständigkeit der Fehlstellensuche fachlich prüfen. DOCX nur Haupttext; EML ohne Anlagen; PDF ohne OCR. Keine externe Übermittlung."}


def report_lines(result):
    yield "# Playbook Prüfung"
    yield f"Review {result['review_id']} · Playbook {result['playbook']['id']} · Version {result['playbook']['version']} · {result['mode']}"
    yield "Der Bericht dokumentiert Regelbefunde und deren Belege. Er ersetzt keine Entscheidung über die Unterzeichnung."
    yield "## 1 Dokumentensatz und Rangfolge"
    yield result["document_precedence"]
    for doc in result["documents"]:
        yield f"{doc['id']}: {doc['title']} · {doc['version']} · {doc['role']} · SHA256 {doc['sha256']}"
    for i, topic in enumerate(result["topics"], 2):
        yield f"## {i} {topic['id']} {topic['title']}"
        yield f"{topic['badge']} · Risikostufe {topic['risk']} · Fundstatus {topic['presence']}"
        yield topic["reasoning"]
        for j, pos in enumerate(topic["positions"], 1):
            yield f"### {i}.{j} {pos['id']} {POSITION_LABELS[pos['type']]}"
            count_label = "Detected" if pos["type"] == "not_acceptable" else "Met"
            yield f"{pos['met']}/{pos['judged']} {count_label} · {pos['not_verifiable']} Not verifiable · {pos['pending']} Pending · {pos['excluded']} Excluded"
            yield f"Verknüpfung: {'alle Bedingungen' if pos['match_mode'] == 'all' else 'mindestens eine Bedingung'}. Position {MATCH_LABELS[pos['match']]}."
            for rule in pos["rules"]:
                yield f"{rule['id']} · {rule['label']} · {rule['condition']}"
                yield rule["reasoning"]
                for cite in rule.get("citations", []):
                    yield f"[{cite['document_id']} · {cite['locator']}] „{cite['quote']}“"
                if rule.get("basis") == "absence":
                    yield "Fehlstellensuche: " + ", ".join(rule["searched_document_ids"]) + ". " + rule["search_description"]
                if rule.get("question"):
                    yield "Rückfrage: " + rule["question"]
                if rule.get("scope_approval"):
                    scope = rule["scope_approval"]
                    yield f"Scope-Ausnahme: {scope['reason']} · {scope['by']} · {scope['reference']}"
        for finding in topic["legal_findings"]:
            yield f"Rechtsbefund {finding['status']}: {finding['reasoning']}"
            yield f"{finding['authority']} · {finding['pinpoint']} · geprüft {finding['checked_on']} · {finding['source_url']}"
            for cite in finding["contract_citations"]:
                yield f"[{cite['document_id']} · {cite['locator']}] „{cite['quote']}“"
    yield f"## {len(result['topics']) + 2} Nächste Entscheidung"
    yield ("Der vollständig befundete Bericht kann zur menschlichen Entscheidung vorgelegt werden." if result["ready_for_human_decision"] else "Vor einer Freigabe sind die markierten Abweichungen oder offenen Punkte zu bearbeiten.")
    yield result["limits"]


def write_docx(result, path):
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise ReviewError("Word-Export benötigt python-docx; keine automatische Installation") from exc
    doc = Document()
    doc.core_properties.title = "Playbook Prüfung"
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
            fonts.attrib.pop(qn("w:" + attr), None)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attr), "Times New Roman")
    for style in doc.styles:
        for borders in list(style.element.iter(qn("w:pBdr"))):
            borders.getparent().remove(borders)
    doc.styles["Title"].font.size = Pt(17)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.4)
    section.top_margin = section.bottom_margin = Cm(2)
    for line in report_lines(result):
        if line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        else:
            doc.add_paragraph(line)
    doc.save(path)


def unique_json(pairs):
    obj = {}
    for key, value in pairs:
        require(key not in obj, f"Doppelter JSON-Schlüssel: {key}")
        obj[key] = value
    return obj


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--source-root", type=Path, help="Standard: Ordner der Eingabe")
    parser.add_argument("--output", type=Path, required=True, help="Neue Zieldatei .json, .md oder .docx")
    parser.add_argument("--draft", action="store_true", help="Pending zulassen; keine Abschlussfreigabe")
    args = parser.parse_args(argv)
    try:
        require(not args.output.exists(), "Zieldatei existiert; neuen Versionsnamen wählen")
        require(args.output.suffix.lower() in {".json", ".md", ".docx"}, "Ausgabeformat muss .json/.md/.docx sein")
        require(args.input.stat().st_size <= 10_000_000, "JSON-Eingabe ist zu groß")
        data = json.loads(args.input.read_text(encoding="utf-8"), object_pairs_hook=unique_json)
        result = evaluate(data, args.source_root or args.input.parent, final=not args.draft)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        if args.output.suffix.lower() == ".docx":
            write_docx(result, args.output)
        elif args.output.suffix.lower() == ".md":
            args.output.write_text("\n\n".join(report_lines(result)) + "\n", encoding="utf-8")
        else:
            args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"Bericht erstellt: {args.output}")
        return 0
    except (ReviewError, OSError, json.JSONDecodeError, KeyError, TypeError, ET.ParseError) as exc:
        print(f"Prüfung abgebrochen: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
